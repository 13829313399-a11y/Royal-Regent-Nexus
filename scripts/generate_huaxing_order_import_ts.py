from __future__ import annotations

import json
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


FIELD_MAP = {
    "machine": "B",
    "isAuto": "D",
    "noteFlag": "E",
    "machineModel": "F",
    "moldCode": "G",
    "productName": "H",
    "orderNo": "I",
    "productCode": "J",
    "quantitySet": "K",
    "orderQty": "L",
    "producedQty": "M",
    "shortageQty": "N",
    "planTarget": "O",
    "color": "Q",
    "colorCode": "R",
    "material": "S",
    "netWeight": "T",
    "grossWeight": "U",
    "materialWeightKg": "V",
    "unitPrice": "W",
    "orderDate": "Z",
    "deliveryStart": "AA",
    "deliveryEnd": "AB",
    "planStart": "AG",
    "planFinish": "AH",
    "planFinishMonth": "AI",
    "inboundDate": "AJ",
    "dueGap": "AK",
    "sprayFlag": "AL",
    "shippingDays": "AM",
    "shiftEndAt": "AN",
    "warehouse": "AR",
    "remark": "AS",
    "shippingDate": "AT",
    "armType": "AU",
    "fixture": "AV",
}


def stringify(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value).strip()


def parse_float(value):
    if value in (None, ""):
        return None
    try:
        return float(value)
    except Exception:
        return None


def excelish_date(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
            return converted.strftime("%Y-%m-%d")
        except Exception:
            return stringify(value)

    return stringify(value)


def derive_customer(order_no: str) -> str:
    if not order_no:
        return "待补客户"

    prefix_match = re.match(r"[A-Z]+", order_no)
    if prefix_match:
        return prefix_match.group(0)

    if "啤机补数" in order_no:
        return "啤机补数"

    return order_no[:4]


def build_issue(row):
    if not row["machine"]:
        return ("待分机", "amber")
    if row["dataAnomaly"]:
        return ("数据异常", "red")
    if row["keyOrder"]:
        return ("关键单", "red")
    if row["overdue"]:
        return ("交期风险", "red")
    if row["remark"] and "转" in row["remark"]:
        return ("转模/转色", "amber")
    if row["remark"] and "复模" in row["remark"]:
        return ("复模排程", "blue")
    return ("正常待排", "blue")


def main() -> int:
    if len(sys.argv) != 3:
        print("usage: generate_huaxing_order_import_ts.py <xlsx-path> <output-ts>", file=sys.stderr)
        return 1

    workbook_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])

    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    rows = []
    material_counts = Counter()
    color_counts = Counter()

    for row_index in range(4, sheet.max_row + 1):
        order_no = stringify(sheet[f"I{row_index}"].value)
        mold_code = stringify(sheet[f"G{row_index}"].value)
        product_name = stringify(sheet[f"H{row_index}"].value)

        if not order_no or not mold_code or not product_name:
            continue

        row = {"sheetRow": row_index}
        for key, column in FIELD_MAP.items():
            raw_value = sheet[f"{column}{row_index}"].value
            if key in {"orderDate", "deliveryStart", "deliveryEnd", "inboundDate"}:
                row[key] = excelish_date(raw_value)
            else:
                row[key] = stringify(raw_value)

        shortage_qty_num = parse_float(sheet[f"N{row_index}"].value)
        order_qty_num = parse_float(sheet[f"L{row_index}"].value)
        produced_qty_num = parse_float(sheet[f"M{row_index}"].value)
        due_gap_num = parse_float(sheet[f"AK{row_index}"].value)

        row["shortageQtyNum"] = shortage_qty_num
        row["orderQtyNum"] = order_qty_num
        row["producedQtyNum"] = produced_qty_num
        row["dueGapNum"] = due_gap_num
        row["pending"] = (shortage_qty_num or 0) > 0
        row["overdue"] = due_gap_num is not None and due_gap_num < 0
        row["keyOrder"] = "关键" in row["noteFlag"]
        row["dataAnomaly"] = shortage_qty_num is not None and shortage_qty_num < 0

        if row["pending"]:
            rows.append(row)
            if row["material"]:
                material_counts[row["material"]] += 1
            if row["color"]:
                color_counts[row["color"]] += 1

    rows.sort(
        key=lambda row: (
            row["dueGapNum"] is None,
            row["dueGapNum"] if row["dueGapNum"] is not None else 999999,
            -(row["shortageQtyNum"] or 0),
        )
    )

    preview_rows = []
    for row in rows[:24]:
        issue, tone = build_issue(row)
        preview_rows.append(
            {
                "orderNo": row["orderNo"],
                "customer": derive_customer(row["orderNo"]),
                "productName": row["productName"],
                "moldCode": row["moldCode"],
                "color": row["color"] or "待补",
                "material": row["material"] or "待补",
                "quantity": f"{int(row['shortageQtyNum'] or 0):,}",
                "dueDate": row["deliveryEnd"] or row["planFinish"] or "待补",
                "cavity": f"套数 {row['quantitySet']}" if row["quantitySet"] else "待补",
                "unitWeight": f"{row['netWeight']} g" if row["netWeight"] else "待补",
                "source": "华兴日排版表 6-30",
                "planner": "日排表导入",
                "machineAdvice": row["machine"] or "待确认",
                "issue": issue,
                "tone": tone,
            }
        )

    summary = {
        "source": str(workbook_path),
        "sheetName": sheet.title,
        "importedAt": datetime.now().strftime("%Y-%m-%d"),
        "pendingOrderCount": len(rows),
        "overdueCount": sum(1 for row in rows if row["overdue"]),
        "missingMachineCount": sum(1 for row in rows if not row["machine"]),
        "dataAnomalyCount": sum(1 for row in rows if row["dataAnomaly"]),
        "keyOrderCount": sum(1 for row in rows if row["keyOrder"]),
        "remarkTransitionCount": sum(1 for row in rows if row["remark"] and "转" in row["remark"]),
        "duplicateOrderCount": len(rows) - len({(row["orderNo"], row["moldCode"], row["machine"], row["planStart"]) for row in rows}),
        "topMaterials": [{"label": label, "count": count} for label, count in material_counts.most_common(8)],
        "topColors": [{"label": label, "count": count} for label, count in color_counts.most_common(8)],
    }

    detail_rows = [
        {
            "sheetRow": row["sheetRow"],
            "machine": row["machine"],
            "isAuto": row["isAuto"],
            "machineModel": row["machineModel"],
            "moldCode": row["moldCode"],
            "productName": row["productName"],
            "orderNo": row["orderNo"],
            "productCode": row["productCode"],
            "shortageQty": row["shortageQty"],
            "planTarget": row["planTarget"],
            "color": row["color"],
            "colorCode": row["colorCode"],
            "material": row["material"],
            "deliveryEnd": row["deliveryEnd"],
            "planStart": row["planStart"],
            "planFinish": row["planFinish"],
            "dueGap": row["dueGap"],
            "remark": row["remark"],
            "armType": row["armType"],
            "fixture": row["fixture"],
            "keyOrder": row["keyOrder"],
            "overdue": row["overdue"],
        }
        for row in rows[:40]
    ]

    content = (
        "export const huaxingOrderImportSummary = "
        + json.dumps(summary, ensure_ascii=False, indent=2)
        + " as const\n\n"
        + "export const huaxingPendingOrderImportRows = "
        + json.dumps(preview_rows, ensure_ascii=False, indent=2)
        + " as const\n\n"
        + "export const huaxingPendingOrderImportDetailRows = "
        + json.dumps(detail_rows, ensure_ascii=False, indent=2)
        + " as const\n"
    )

    output_path.write_text(content, encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
