from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def main() -> int:
    if len(sys.argv) < 3:
        print("usage: dump_sheet_rows.py <xlsx-path> <row> [<row>...]", file=sys.stderr)
        return 1

    workbook_path = Path(sys.argv[1])
    row_indexes = [int(value) for value in sys.argv[2:]]

    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    payload = []
    for row_index in row_indexes:
        cells = []
        for col_index in range(1, min(sheet.max_column, 220) + 1):
            value = sheet.cell(row=row_index, column=col_index).value
            if value is None or value == "":
                continue
            cells.append(
                {
                    "cell": sheet.cell(row=row_index, column=col_index).coordinate,
                    "value": str(value),
                }
            )
        payload.append({"row": row_index, "cells": cells})

    print(json.dumps(payload, ensure_ascii=True, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
