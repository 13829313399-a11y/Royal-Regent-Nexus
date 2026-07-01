from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def main() -> int:
    if len(sys.argv) < 2 or len(sys.argv) > 4:
        print("usage: inspect_sheet_layout.py <xlsx-path> [max-rows] [max-cols]", file=sys.stderr)
        return 1

    workbook_path = Path(sys.argv[1])
    max_rows = int(sys.argv[2]) if len(sys.argv) >= 3 else 12
    max_cols = int(sys.argv[3]) if len(sys.argv) >= 4 else 80

    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    preview = []
    for row_index in range(1, max_rows + 1):
        row_cells = []
        for col_index in range(1, max_cols + 1):
            value = sheet.cell(row=row_index, column=col_index).value
            if value is None or value == "":
                continue
            row_cells.append(
                {
                    "cell": sheet.cell(row=row_index, column=col_index).coordinate,
                    "value": str(value),
                }
            )
        preview.append({"row": row_index, "cells": row_cells})

    merged_ranges = [str(item) for item in list(sheet.merged_cells.ranges)[:40]]

    payload = {
        "path": str(workbook_path),
        "sheetName": sheet.title,
        "maxRow": sheet.max_row,
        "maxColumn": sheet.max_column,
        "mergedRanges": merged_ranges,
        "preview": preview,
    }

    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
