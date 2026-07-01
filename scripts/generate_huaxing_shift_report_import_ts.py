from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


FIELD_MAP = {
    "machineCode": "A",
    "shiftCode": "B",
    "worker": "C",
    "orderNo": "D",
    "moldCode": "E",
    "actualOutput": "F",
    "downtimeMinutes": "G",
    "downtimeReason": "H",
    "carryOverQty": "I",
    "deliveryCode": "J",
    "inboundQty": "K",
    "remark": "L",
}


def stringify(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value).strip()


def parse_number(value):
    if value in (None, ""):
        return 0
    try:
        return float(str(value).replace(",", "").strip())
    except Exception:
        return 0


def build_tone(actual_output: float, carry_over_qty: float, downtime_minutes: float, downtime_reason: str):
    if carry_over_qty > 0 or downtime_minutes >= 60:
        return "red"
    if downtime_minutes > 0 or downtime_reason:
        return "amber"
    if actual_output > 0:
        return "green"
    return "blue"


def main() -> int:
    if len(sys.argv) != 3:
        print("usage: generate_huaxing_shift_report_import_ts.py <xlsx-path> <output-ts>", file=sys.stderr)
        return 1

    workbook_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])

    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    rows = []

    for row_index in range(2, sheet.max_row + 1):
        machine_code = stringify(sheet[f"{FIELD_MAP['machineCode']}{row_index}"].value)
        shift_code = stringify(sheet[f"{FIELD_MAP['shiftCode']}{row_index}"].value)
        order_no = stringify(sheet[f"{FIELD_MAP['orderNo']}{row_index}"].value)

        if not machine_code or not shift_code or not order_no:
            continue

        row = {"sheetRow": row_index}
        for key, column in FIELD_MAP.items():
            row[key] = stringify(sheet[f"{column}{row_index}"].value)

        actual_output_num = parse_number(row["actualOutput"])
        downtime_minutes_num = parse_number(row["downtimeMinutes"])
        carry_over_qty_num = parse_number(row["carryOverQty"])
        inbound_qty_num = parse_number(row["inboundQty"])
        tone = build_tone(actual_output_num, carry_over_qty_num, downtime_minutes_num, row["downtimeReason"])

        row["actualOutputNum"] = actual_output_num
        row["downtimeMinutesNum"] = downtime_minutes_num
        row["carryOverQtyNum"] = carry_over_qty_num
        row["inboundQtyNum"] = inbound_qty_num
        row["tone"] = tone

        rows.append(row)

    summary = {
        "source": str(workbook_path),
        "sheetName": sheet.title,
        "importedAt": datetime.now().strftime("%Y-%m-%d"),
        "reportCount": len(rows),
        "handoverCount": sum(1 for row in rows if row["carryOverQtyNum"] > 0),
        "downtimeCount": sum(1 for row in rows if row["downtimeMinutesNum"] > 0 or row["downtimeReason"]),
        "inboundLinkedCount": sum(1 for row in rows if row["deliveryCode"] or row["inboundQtyNum"] > 0),
    }

    detail_rows = [
        {
            "sheetRow": row["sheetRow"],
            "machineCode": row["machineCode"],
            "shiftCode": row["shiftCode"],
            "worker": row["worker"],
            "orderNo": row["orderNo"],
            "moldCode": row["moldCode"],
            "actualOutput": row["actualOutput"],
            "downtimeMinutes": row["downtimeMinutes"],
            "downtimeReason": row["downtimeReason"],
            "carryOverQty": row["carryOverQty"],
            "deliveryCode": row["deliveryCode"],
            "inboundQty": row["inboundQty"],
            "remark": row["remark"],
            "tone": row["tone"],
        }
        for row in rows
    ]

    content = (
        "export const huaxingShiftReportImportSummary = "
        + json.dumps(summary, ensure_ascii=False, indent=2)
        + " as const\n\n"
        + "export const huaxingShiftReportImportRows = "
        + json.dumps(detail_rows, ensure_ascii=False, indent=2)
        + " as const\n"
    )

    output_path.write_text(content, encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
