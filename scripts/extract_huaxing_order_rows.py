from __future__ import annotations

import json
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


FIELD_MAP = {
    "machine": "B",
    "is_auto": "D",
    "note_flag": "E",
    "machine_model": "F",
    "mold_code": "G",
    "product_name": "H",
    "order_no": "I",
    "product_code": "J",
    "quantity_set": "K",
    "order_qty": "L",
    "produced_qty": "M",
    "shortage_qty": "N",
    "plan_target": "O",
    "color": "Q",
    "color_code": "R",
    "material": "S",
    "net_weight": "T",
    "gross_weight": "U",
    "material_weight_kg": "V",
    "unit_price": "W",
    "order_date": "Z",
    "delivery_start": "AA",
    "delivery_end": "AB",
    "plan_start": "AG",
    "plan_finish": "AH",
    "plan_finish_month": "AI",
    "inbound_date": "AJ",
    "due_gap": "AK",
    "warehouse": "AR",
    "remark": "AS",
    "shipping_date": "AT",
    "arm_type": "AU",
    "fixture": "AV",
}


def stringify(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    return str(value)


def parse_float(value):
    if value is None or value == "":
        return None
    try:
        return float(value)
    except Exception:
        return None


def main() -> int:
    if len(sys.argv) != 2:
        print("usage: extract_huaxing_order_rows.py <xlsx-path>", file=sys.stderr)
        return 1

    workbook_path = Path(sys.argv[1])
    workbook = load_workbook(workbook_path, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    rows = []
    customer_prefixes = Counter()
    material_counts = Counter()
    color_counts = Counter()
    machine_counts = Counter()

    for row_index in range(4, sheet.max_row + 1):
        order_no = stringify(sheet[f"I{row_index}"].value).strip()
        mold_code = stringify(sheet[f"G{row_index}"].value).strip()
        product_name = stringify(sheet[f"H{row_index}"].value).strip()

        if not order_no or not mold_code or not product_name:
            continue

        shortage_qty = parse_float(sheet[f"N{row_index}"].value)
        order_qty = parse_float(sheet[f"L{row_index}"].value)
        produced_qty = parse_float(sheet[f"M{row_index}"].value)
        due_gap = parse_float(sheet[f"AK{row_index}"].value)

        row = {"sheet_row": row_index}
        for key, column in FIELD_MAP.items():
            row[key] = stringify(sheet[f"{column}{row_index}"].value).strip()

        row["shortage_qty_num"] = shortage_qty
        row["order_qty_num"] = order_qty
        row["produced_qty_num"] = produced_qty
        row["due_gap_num"] = due_gap
        rows.append(row)

        customer_prefixes[order_no[:3]] += 1
        if row["material"]:
            material_counts[row["material"]] += 1
        if row["color"]:
            color_counts[row["color"]] += 1
        if row["machine"]:
            machine_counts[row["machine"]] += 1

    positive_shortage_rows = [row for row in rows if (row["shortage_qty_num"] or 0) > 0]

    payload = {
        "path": str(workbook_path),
        "sheetName": sheet.title,
        "rowCount": len(rows),
        "positiveShortageCount": len(positive_shortage_rows),
        "customerPrefixes": customer_prefixes.most_common(10),
        "topMaterials": material_counts.most_common(10),
        "topColors": color_counts.most_common(10),
        "topMachines": machine_counts.most_common(10),
        "previewRows": positive_shortage_rows[:20],
    }

    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
