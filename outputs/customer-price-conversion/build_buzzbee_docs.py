from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("outputs/customer-price-conversion")
REQ_MD = OUT_DIR / "客价转换台-BuzzBee内部转报客需求文档.md"
FORM_DOCX = OUT_DIR / "客价转换规则收集表-跟客填写模板.docx"


REQUIREMENT_MD = """# 客价转换台 - BuzzBee 内部转报客需求文档

## 1. 背景与目标

业务部「客价转换台」需要把各跟客上传的内部报价 Excel 转换为对应客户的报客价 Excel。权限按车间跟客账号管理，不按单个客户拆账号；同一个车间跟客账号可以处理该车间绑定的多个客户。BuzzBee 是第一套要落地的客户模板。

本需求基于旧版「内部转报客网页」目录中的源码、使用说明、内部价样例和报客价样例整理，目标是把旧工具的可用逻辑沉淀到新系统中，并为后续其他客户模板提供统一采集口径。

## 2. 资料来源

- 旧网页入口：`index.html`
- 默认参数：`js/defaults.js`
- 内部 Excel 解析：`js/parser.js`
- 报客 Excel 生成：`js/exporter.js`
- 报客 Excel 反向解析：`js/parse-client.js`
- 报客价对比：`js/compare.js`
- 样例内部价：`露营火堆套装枪报价2026-3-15（内部价钱）.xlsx`
- 样例报客价：`R0 RR ltem 露营火堆套装 box(2026-3-18）（报客价钱）.xls`

## 3. 用户流程

1. 跟客登录系统，进入「业务部 / 报价与成本中心 / 客价转换台」。
2. 选择本人车间账号下绑定的客户，例如 BuzzBee。
3. 上传该客户的内部报价 Excel。
4. 系统识别一个或多个产品 Sheet，解析注塑件、外购件、工艺人工、纸箱和彩盒资料。
5. 系统按客户模板规则计算报客价，并生成报客价 Excel。
6. 下方展示多 Sheet / 多版本明细对比，可比较历史导出版本或外部上传的报客价文件。

## 4. 权限规则

- 账号粒度：按车间跟客账号分配，不按客户单独开户。
- 客户范围：跟客账号只看到自己账号绑定的客户集合。
- BuzzBee 当前归属：`huaxing_molding_a_sales`，登录页显示「华兴跟客」。
- 后续扩展：新增客户时，只需在客户绑定表中维护客户与车间跟客账号的关系。
- 管理员/业务经理：后续可拥有跨车间查看、复核、客户分配和模板维护权限。

## 5. BuzzBee 内部报价识别规则

### 5.1 Sheet 选择

- 优先读取名称包含「明细」且存在注塑件表头的 Sheet。
- 表头判断：C 列为「名称」，D 列为「料型」。
- 如果没有「明细」Sheet，则回退到第一个包含上述表头的 Sheet。
- 如果仍找不到，提示：未找到含注塑件明细的 Sheet。

样例确认：

- 内部价样例 Sheet：`明细`
- 使用范围：`A1:U80`
- 注塑件表头：第 8 行
- 成本分解起点：第 25 行
- 彩盒信息锚点：第 44 行附近

### 5.2 产品名识别

- 在前 12 行查找包含「报价」「套装」「按图」的文本。
- 去掉括号内容后作为产品名。
- 如果找不到，用 Sheet 名去掉末尾「明细」作为产品名。

### 5.3 注塑件字段

从表头下一行开始逐行读取，遇到合计/小计/空段结束。

| 内部列 | 字段 | 用途 |
|---|---|---|
| C | 名称 | 报客 DESCRIPTION |
| D | 料型 | 映射英文料型与报客单价 |
| E | 克重(G) | 计算材料金额 |
| G | 机型/吨位 | 查机型工费 |
| H | 1啤套数 | 计算啤工 |
| I | 目标产量 | 计算啤工 |
| J | 啤工 | 旧表内部值，仅参考 |
| K | 料金额 | 旧表内部值，仅参考 |

### 5.4 成本分解字段

- 从 B 列出现「料价」的行开始扫描。
- 读取 A 列税别标签、B 列分类、C 列描述、D 列内部金额。
- 跳过报客价、目标价、相差、旺季价、含税、利润、总成本等分析/汇总行。
- 连续空白达到阈值后停止。

### 5.5 纸箱和彩盒

- 纸箱尺寸：查找「外箱:」「外箱：」「外箱」后读取长、宽、高。
- 装箱数：查找「装箱:」「装箱：」「装箱」后读取 PCS/CARTON。
- 纸箱内部价：查找「合计」旁边的数字。
- CUFT：查找「CUFT:」；如缺失，按 `长 * 宽 * 高 / 1728` 计算。
- 彩盒：查找「报客彩盒」，读取下两行的报客彩盒价、FSC 价和 MOQ。若 FSC 缺失且有彩盒价，默认按 `彩盒价 * 1.03` 补值。

## 6. 默认参数

### 6.1 料型参数

料型使用「内部料型名 -> 英文名 -> 报客单价(HKD/KG)」映射。料价不再用内部单价乘倍率，而是直接使用报客单价。

| 内部料型 | 英文名 | 默认报客单价 HKD/KG |
|---|---:|---:|
| PP料 | PP | 12.1 |
| ABS料 | ABS | 16.5 |
| PC料 | PC | 18.0 |
| K料 | K | 22.0 |
| GP料 | GP | 11.0 |
| PE料 | PE | 14.5 |
| POM料 | POM | 28.0 |
| 475料 | 475 | 12.0 |
| C-ABS料 | C-ABS | 31.3 |
| EVA料 | EVA | 13.0 |
| 马力士 | Mars | 14.0 |
| 环保PVC | PVC | 14.0 |
| 尼龙料 | Nylon | 34.0 |
| TPR料 | TPR | 28.0 |
| SAN料 | SAN | 20.0 |
| PP料明 | PP(C) | 18.0 |
| HI75%+GP25% | HI/GP | 12.0 |
| HDPE料 | HDPE | 13.0 |

### 6.2 机型工费

按吨位上限查找第一条满足 `吨位 <= 上限` 的机型工费。

| 吨位上限 | 工费 元/11h |
|---:|---:|
| 6 | 1040 |
| 9 | 1160 |
| 12 | 1280 |
| 14 | 1650 |
| 18 | 1890 |
| 24 | 2130 |
| 34 | 2460 |
| 44 | 2700 |
| 49.9 | 2790 |
| 999 | 3090 |

### 6.3 加价规则

| 成本类型 | 报客计算规则 |
|---|---|
| 料价 | 使用料型参数中的报客单价 |
| 啤工 | 计算内部啤工后乘 1.10 |
| 装配工 | 内部价乘 1.00 |
| 油漆 | 内部价乘 1.00 |
| 喷油工 | 内部价乘 1.00 |
| 纸箱 | 内部价乘 1.03 |
| 车衣 | 内部价乘 1.03 |
| 其它外购 | 内部价乘 1.05 |
| 运费 | 不计入报客转换 |

## 7. 计算逻辑

### 7.1 注塑件

- `PRICE/KG = 料型报客单价`
- `AMOUNT = 克重(G) * PRICE/KG / 1000`
- `机型工费 = 按机型/吨位查表`
- `啤工内部 = 机型工费 / 目标产量 / 1啤套数`
- `啤工 COST = 啤工内部 * 啤工加价倍率`

### 7.2 外购件

- 跳过分类：料价、啤工、纸箱、彩盒、杂项。
- 装配工、油漆、喷油工不进入外购明细，汇总到 PROCESS。
- 描述中如包含 `(nPCS)`，解析为数量 n，并从描述中移除该数量标签。
- `单价 = 内部金额 / 数量 * 分类加价倍率`
- `金额 = 单价 * 数量`
- 描述包含「子弹」「波波球」「波波」「加值件」「附加件」时，放入 Additional Parts。

### 7.3 汇总

- `PLASTIC = 注塑件 AMOUNT 合计`
- `INJECTION = 注塑件啤工 COST 合计`
- `PURCHASE = 外购件 AMOUNT 合计`
- `PROCESS = 油漆 + 喷油工 + 装配工`
- `SUB = PLASTIC + INJECTION + PURCHASE + PROCESS`
- `P+O = SUB * 10%`
- `BREAKDOWN.TOTAL = SUB + P+O`
- `TOTAL Ex-fty cost = BREAKDOWN.TOTAL + Additional Parts 合计`
- `US$ = TOTAL Ex-fty cost / 7.75`

## 8. BuzzBee 报客 Excel 输出格式

输出为一个工作簿；如果内部报价含多个产品 Sheet，则每个产品输出一个报客 Sheet。

关键布局：

- A1：`COST BREAKDOWN SHEET (ROYAL REGENT)`
- F1:G2：DATE / REVISED
- A4：`ltem:` + 产品名
- A5：`INJECTION`
- A6:G6：注塑件表头 DESCRIPTION / MAT'L / WEIGHT / PRICE/KG / AMOUNT / SET/SHOT / COST
- A7 起：注塑件明细，至少保留 17 行注塑区
- 注塑区后：SUB TOTAL
- PURCHASE 区：DESCRIPTION / QTY / PRICE / AMOUNT
- 右侧 BREAKDOWN 区：PLASTIC / INJECTION / PURCHASE / TRAN / PROCESS / SUB / P+O / TOTAL
- Additional Parts 区：DESCRIPTION / QTY / PRICE / 金额
- TOTAL Ex-fty cost、US$
- CARTON SIZE、PACK/CUBE、OUTTER
- 彩盒价格、报客彩盒 FSC、MOQ

样例报客表确认：

- Sheet：`露营火堆套装`
- 使用范围：`A1:J58`
- INJECTION：第 5 行
- PURCHASE：第 26 行
- Additional Parts：第 40 行
- TOTAL Ex-fty cost：第 43 行
- US$：第 45 行
- CARTON SIZE：第 55 行
- OUTTER：第 58 行

## 9. 报客价对比逻辑

对比模块用于比较两份报客价或两个导出版本。

### 9.1 反向解析锚点

- 必须包含 `INJECTION` 段才认为是有效报客价。
- 产品名读取 A4 的 `ltem:`。
- 注塑件从 INJECTION 表头后读取至 SUB TOTAL。
- 外购件从 PURCHASE 表头后读取至 SUB TOTAL。
- BREAKDOWN 从 I/J 列读取 PLASTIC、INJECTION、PURCHASE、TRAN、PROCESS、SUB、P+O、TOTAL。
- Additional Parts 从对应标签后读取至 TOTAL Ex-fty cost 前。
- 纸箱尺寸读取 CARTON SIZE / OUTTER。
- 彩盒价格读取 CARTON SIZE 附近 F/G/H 列。

### 9.2 对比内容

- 大项汇总：PLASTIC、INJECTION、PURCHASE、PROCESS、SUB、P+O、BREAKDOWN.TOTAL、TOTAL Ex-fty cost、US$。
- 工艺人工：SPRAY、ASSEMBLY、PROCESS 总。
- 注塑件：按名称匹配，比较 qty、weight、pricePerKg、setShot、price、amount、cost。
- 外购件：按描述匹配，识别缺项和新增项。
- Additional Parts：按描述匹配。
- 纸箱：LENGTH、WIDTH、HEIGHT、PCS/CARTON、CUBE。
- 彩盒：报客彩盒 1/2、FSC、MOQ。

## 10. 新系统实现建议

### 10.1 数据模型建议

- 客户模板：客户、模板版本、状态、适用日期。
- 跟客客户绑定：跟客账号、客户、车间、是否启用。
- 料型映射：客户模板、内部料型名、英文名、内部单价、报客单价。
- 机型工费：客户模板、吨位上限、工费。
- 加价规则：客户模板、成本分类、倍率、是否计入外购/PROCESS/Additional Parts。
- 导入批次：客户、跟客账号、源文件名、Sheet 数、解析状态。
- 导出版本：客户、导入批次、导出文件名、版本号、汇总金额。
- 明细行：产品 Sheet、注塑件、外购件、Additional Parts、纸箱、彩盒、汇总项。

### 10.2 系统功能

- 支持客户模板选择与版本管理。
- 支持上传内部报价 Excel 并显示解析结果。
- 支持异常提示：找不到表头、料型未映射、机型未命中、纸箱/彩盒缺字段。
- 支持导出客户格式 Excel。
- 支持导出版本保留与明细对比。
- 支持参数维护：料型、机型、倍率、Additional Parts 关键字、汇率/美元换算。
- 支持每个客户独立模板，但权限仍按车间跟客账号绑定客户集合。

## 11. 风险与待确认

- BuzzBee 样例中部分中文显示在旧环境里依赖编码处理，新系统应使用统一 UTF-8 / xlsx 解析方式。
- `.xls` 报客样例为老式二进制格式，新系统建议导出 `.xlsx`，如客户要求 `.xls` 再单独适配。
- 料型报客单价属于客户模板核心参数，需要业务确认维护方式和审批权限。
- 汇总中的 `US$ = Ex-fty / 7.75` 是否固定，需要业务确认是否做成客户参数。
- `P+O = 10%` 是否固定，需要业务确认。
- Additional Parts 当前按关键词识别，其他客户可能需要独立规则。
- 彩盒 FSC 缺失时默认 1.03 倍，需确认是否所有客户适用。

## 12. 后续其他客户采集方式

给每个跟客一份《客价转换规则收集表》，要求填写客户格式差异，并同时附：

- 该客户真实内部报价 Excel。
- 该客户真实报客价 Excel。
- 如有多个产品或多个 Sheet，至少提供一个多 Sheet 样例。
- 如有特殊料型、特殊加价、特殊排版，需在 Word 表中写明。
"""


def set_font(run, name="Calibri", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_style(paragraph, after=6, before=0, line=1.25):
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(after)
    fmt.space_before = Pt(before)
    fmt.line_spacing = line


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                set_paragraph_style(p, after=0, line=1.15)
                for run in p.runs:
                    set_font(run, size=9)


def add_heading(doc, text, level):
    paragraph = doc.add_heading(level=level)
    run = paragraph.add_run(text)
    for old_run in paragraph.runs[:-1]:
        old_run.text = ""
    set_font(run, size=16 if level == 1 else 13 if level == 2 else 12, bold=True, color="2E74B5" if level < 3 else "1F4D78")
    set_paragraph_style(paragraph, before=18 if level == 1 else 14 if level == 2 else 10, after=8 if level == 1 else 6)
    return paragraph


def add_body(doc, text):
    p = doc.add_paragraph()
    set_paragraph_style(p)
    run = p.add_run(text)
    set_font(run, size=11)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_paragraph_style(p, after=4)
        run = p.add_run(item)
        set_font(run, size=10.5)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        set_cell_shading(cell, "E8EEF5")
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_font(run, size=9, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    set_table_geometry(table, widths)
    doc.add_paragraph()
    return table


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_style(title, after=3, line=1.15)
    run = title.add_run("客价转换规则收集表")
    set_font(run, size=22, bold=True, color="0B2545")

    subtitle = doc.add_paragraph()
    set_paragraph_style(subtitle, after=10)
    run = subtitle.add_run("给各车间跟客填写：请连同客户内部报价 Excel、报客价 Excel 一并提交，用于新增客户模板。")
    set_font(run, size=10.5, color="555555")

    add_heading(doc, "1. 填写人和客户信息", 1)
    add_table(
        doc,
        ["项目", "填写内容"],
        [
            ["客户名称", ""],
            ["所属工厂 / 车间", ""],
            ["跟客账号 / 姓名", ""],
            ["客户报价格式负责人", ""],
            ["样例产品名称", ""],
            ["提交日期", ""],
        ],
        [2700, 6660],
    )

    add_heading(doc, "2. 必须随表附上的文件", 1)
    add_table(
        doc,
        ["附件", "是否提供", "说明"],
        [
            ["内部报价 Excel", "□ 是  □ 否", "请提供原始文件，不要截图。"],
            ["报客价 Excel", "□ 是  □ 否", "请提供已发客户或准备发客户的版本。"],
            ["多 Sheet 样例", "□ 是  □ 否", "若客户报价常有多个产品 Sheet，请至少给一份。"],
            ["特殊规则说明", "□ 是  □ 否", "如有额外邮件、手工改价、固定折扣，请写在第 9 节。"],
        ],
        [2300, 1900, 5160],
    )

    add_heading(doc, "3. 内部报价 Excel 识别规则", 1)
    add_body(doc, "请按真实客户文件填写。系统后续会根据这些锚点识别内部报价，不同客户可以有不同格式。")
    add_table(
        doc,
        ["问题", "填写内容"],
        [
            ["有效 Sheet 命名规则", "例如：名称包含“明细”；或固定第几个 Sheet。"],
            ["产品名所在位置", "例如：前 10 行某个标题；或 Sheet 名。"],
            ["注塑件表头位置", "例如：C列=名称，D列=料型。"],
            ["注塑件明细结束规则", "例如：遇到合计/空行/小计停止。"],
            ["外购件/成本分解起点", "例如：B列出现“料价”。"],
            ["纸箱尺寸位置", "例如：外箱、装箱、CUFT 标签。"],
            ["彩盒价格位置", "例如：报客彩盒下两行。"],
        ],
        [3100, 6260],
    )

    add_heading(doc, "4. 报客价 Excel 输出格式", 1)
    add_table(
        doc,
        ["区域", "客户要求 / 单元格位置"],
        [
            ["标题 / 日期 / Revision", ""],
            ["产品名 / Item", ""],
            ["注塑件区域", ""],
            ["外购件区域", ""],
            ["BREAKDOWN 汇总", ""],
            ["Additional Parts", ""],
            ["TOTAL Ex-fty cost / US$", ""],
            ["纸箱 / OUTTER / CUBE", ""],
            ["彩盒 / FSC / MOQ", ""],
        ],
        [3100, 6260],
    )

    add_heading(doc, "5. 料型和材料单价规则", 1)
    add_body(doc, "请列出客户实际报客用的料型名称和单价。若客户报价不按 HKD/KG，请写明单位。")
    add_table(
        doc,
        ["内部料型写法", "报客显示名称", "报客单价", "单位", "备注"],
        [["", "", "", "", ""] for _ in range(8)],
        [2100, 2100, 1500, 1200, 2460],
    )

    add_heading(doc, "6. 加价和计算规则", 1)
    add_table(
        doc,
        ["成本类型", "客户规则", "是否进入明细", "是否进入汇总"],
        [
            ["料价", "", "□ 是  □ 否", ""],
            ["啤工", "", "□ 是  □ 否", ""],
            ["装配工", "", "□ 是  □ 否", ""],
            ["油漆", "", "□ 是  □ 否", ""],
            ["喷油工", "", "□ 是  □ 否", ""],
            ["纸箱", "", "□ 是  □ 否", ""],
            ["车衣", "", "□ 是  □ 否", ""],
            ["其它外购", "", "□ 是  □ 否", ""],
            ["运费", "", "□ 是  □ 否", ""],
        ],
        [1800, 3960, 1800, 1800],
    )

    add_heading(doc, "7. 汇总公式", 1)
    add_table(
        doc,
        ["项目", "客户公式 / 说明"],
        [
            ["PLASTIC", ""],
            ["INJECTION", ""],
            ["PURCHASE", ""],
            ["PROCESS", ""],
            ["SUB", ""],
            ["P+O / 利润 / 管理费", ""],
            ["TOTAL / Ex-fty", ""],
            ["US$ / 汇率", ""],
        ],
        [2700, 6660],
    )

    add_heading(doc, "8. 特殊项目和关键字", 1)
    add_body(doc, "例如 BuzzBee 旧规则会把“子弹、波波球、波波、加值件、附加件”放入 Additional Parts。其他客户如有类似特殊区，请在这里写清楚。")
    add_table(
        doc,
        ["关键字 / 分类", "处理方式", "示例"],
        [["", "", ""], ["", "", ""], ["", "", ""], ["", "", ""]],
        [2600, 3900, 2860],
    )

    add_heading(doc, "9. 人工调整和例外情况", 1)
    add_body(doc, "请写明报客价是否存在手工改价、固定小数位、四舍五入、最低价、客户指定隐藏项目、客户指定排序等规则。")
    add_table(
        doc,
        ["例外项", "说明"],
        [["", ""], ["", ""], ["", ""], ["", ""]],
        [2600, 6760],
    )

    add_heading(doc, "10. 跟客确认", 1)
    add_table(
        doc,
        ["确认项", "结果"],
        [
            ["我已提供内部报价 Excel 和报客价 Excel 原文件", "□ 是  □ 否"],
            ["我已确认以上规则能代表该客户常用报价格式", "□ 是  □ 否"],
            ["如系统转换结果与手工表不同，我愿意协助核对差异", "□ 是  □ 否"],
            ["签名 / 日期", ""],
        ],
        [5600, 3760],
    )

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = footer.add_run("客价转换台规则收集模板")
        set_font(run, size=9, color="777777")

    doc.save(FORM_DOCX)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    REQ_MD.write_text(REQUIREMENT_MD, encoding="utf-8")
    build_docx()
    print(REQ_MD.resolve())
    print(FORM_DOCX.resolve())


if __name__ == "__main__":
    main()
