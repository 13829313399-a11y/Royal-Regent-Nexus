import importlib
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font
from PIL import Image

from test_internal_quote_api import make_client, login, logout, create_payload
from test_internal_quote_alternatives import fill, fork, create


def image_bytes():
    buffer = BytesIO()
    Image.new("RGB", (8, 8), "green").save(buffer, format="PNG")
    return buffer.getvalue()


def series(client, suffix="SERIES"):
    payload = create_payload(suffix=suffix)
    payload.update(workflow_mode="direct_output", quote_type="series", products=[
        {"product_name": name, "qty": 5000} for name in ("彩色兔子", "粉色兔子", "黄色兔子")
    ])
    response = client.post("/api/internal-quotes", json=payload)
    assert response.status_code == 201, response.text
    products = client.get(f"/api/internal-quotes/{response.json()['id']}/batch-products").json()
    return [client.get(f"/api/internal-quotes/{row['quote_id']}").json() for row in products]


def request_body(quotes):
    return {"products": [{"quote_id": row["id"], "revision": row["header_revision"]} for row in quotes]}


def test_series_workbook_preserves_images_styles_printing_and_isolated_formula_dependencies():
    from app.services.internal_quote_series_export import combine_series_workbooks
    sources = []
    for amount in (10, 20, 30):
        book = Workbook()
        sheet = book.active
        sheet.title = "报价明细"
        sheet["A1"] = "='电子明细'!B2"
        sheet["A2"] = '=IF(A1>0,"电子明细!B2",0)'
        sheet["B3"] = amount
        sheet["B3"].font = Font(bold=True, color="FF008800")
        sheet.merge_cells("C3:D3")
        sheet.row_dimensions[3].height = 35
        sheet.print_area = "A1:D10"
        sheet.print_title_rows = "1:2"
        sheet.add_image(ExcelImage(BytesIO(image_bytes())), "C5")
        data = book.create_sheet("电子明细")
        data["B2"] = amount
        data.sheet_state = "veryHidden"
        buffer = BytesIO(); book.save(buffer); book.close()
        sources.append(("同名/产品" * 9, buffer.getvalue()))
    output = load_workbook(BytesIO(combine_series_workbooks(sources)))
    visible = [sheet for sheet in output if sheet.sheet_state == "visible"]
    assert len(visible) == 3
    assert len({sheet.title for sheet in visible}) == 3
    for index, sheet in enumerate(visible, 1):
        assert len(sheet.title) <= 31 and "/" not in sheet.title
        assert sheet["A1"].value == f"='{index}-电子明细'!B2"
        assert sheet["A2"].value == '=IF(A1>0,"电子明细!B2",0)'
        assert output[f"{index}-电子明细"]["B2"].value == index * 10
        assert sheet["B3"].font.bold and sheet.row_dimensions[3].height == 35
        assert len(sheet._images) == 1 and "C3:D3" in sheet.merged_cells
        assert "$A$1:$D$10" in sheet.print_area and sheet.print_title_rows == "$1:$2"
    output.close()


def test_series_issue_is_atomic_reuses_frozen_exports_and_accepts_one_alternative_per_product(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "series-owner", "admin", "*", "*")
        quotes = series(client)
        quotes[0] = fill(client, quotes[0])
        endpoint = f"/api/internal-quotes/{quotes[0]['id']}/series-export"
        denied = client.post(endpoint, json=request_body(quotes))
        assert denied.status_code == 409 and "粉色兔子" in denied.text
        for quote in quotes:
            current = client.get(f"/api/internal-quotes/{quote['id']}").json()
            assert current["final_release_status"] != "issued"
            assert client.get(f"/api/internal-quotes/{quote['id']}/exports").json() == []
        quotes[1:] = [fill(client, quote) for quote in quotes[1:]]
        service = importlib.import_module("app.services.internal_quote_series_export")
        with monkeypatch.context() as failure:
            def fail_combination(_products):
                raise ValueError("simulated workbook failure")
            failure.setattr(service, "combine_series_workbooks", fail_combination)
            with pytest.raises(ValueError, match="simulated workbook failure"):
                client.post(endpoint, json=request_body(quotes))
        for quote in quotes:
            assert client.get(f"/api/internal-quotes/{quote['id']}/exports").json() == []
        first = client.post(f"/api/internal-quotes/{quotes[0]['id']}/direct-issue", json={"revision": quotes[0]["header_revision"]}).json()
        frozen_url = f"/api/internal-quotes/{quotes[0]['id']}/exports/{first['id']}/download"
        frozen = client.get(frozen_url).content
        alternative = fork(client, quotes[1], "version")
        selected = [quotes[0], alternative, quotes[2]]
        response = client.post(endpoint, json=request_body(selected))
        assert response.status_code == 200, response.text[:1000]
        workbook = load_workbook(BytesIO(response.content))
        assert [sheet.title for sheet in workbook if sheet.sheet_state == "visible"] == [row["product_name"] for row in quotes]
        workbook.close()
        assert client.get(frozen_url).content == frozen
        assert client.get(f"/api/internal-quotes/{quotes[1]['id']}").json()["final_release_status"] != "issued"
        repeated = client.post(endpoint, json=request_body(selected))
        assert repeated.status_code == 200
        for quote in selected:
            exports = client.get(f"/api/internal-quotes/{quote['id']}/exports").json()
            assert len(exports) == 1


def test_series_rejects_missing_duplicate_foreign_and_stale_products(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "series-validation", "admin", "*", "*")
        quotes = series(client)
        endpoint = f"/api/internal-quotes/{quotes[0]['id']}/series-export"
        other = create(client, "OTHER-SERIES")
        for members in (quotes[:2], [quotes[0], quotes[0], quotes[2]], [quotes[0], quotes[1], other]):
            assert client.post(endpoint, json=request_body(members)).status_code == 400
        stale = request_body(quotes)
        stale["products"][0]["revision"] += 1
        assert client.post(endpoint, json=stale).status_code == 409
        logout(client)
        login(client, "series-no-export", "engineer", "engineering")
        assert client.post(endpoint, json=request_body(quotes)).status_code == 403


def test_product_image_in_quote_preserves_issued_version_and_rejects_stale_or_invalid_upload(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "main-image-owner", "admin", "*", "*")
        quote = fill(client, create(client, "MAIN-IMAGE"))
        url = f"/api/internal-quotes/{quote['id']}/product-image"
        def upload(url, revision, content=None):
            return client.post(url, data={"revision": revision}, files={"file": ("main.png", content if content is not None else image_bytes(), "image/png")})
        assert upload(url, quote["header_revision"] + 1).status_code == 409
        assert upload(url, quote["header_revision"], b'\x89PNG\r\n\x1a\ninvalid').status_code == 400
        response = upload(url, quote["header_revision"])
        assert response.status_code == 201, response.text
        quote = client.get(f"/api/internal-quotes/{quote['id']}").json()
        issued = client.post(f"/api/internal-quotes/{quote['id']}/direct-issue", json={"revision": quote["header_revision"]}).json()
        quote = client.get(f"/api/internal-quotes/{quote['id']}").json()
        download = f"/api/internal-quotes/{quote['id']}/exports/{issued['id']}/download"
        frozen = client.get(download).content
        assert upload(url, quote["header_revision"]).status_code == 409
        new = fork(client, quote, "version")
        picture = BytesIO(); Image.new("RGB", (10, 10), "red").save(picture, format="PNG")
        assert upload(f"/api/internal-quotes/{new['id']}/product-image", new["header_revision"], picture.getvalue()).status_code == 201
        assert client.get(download).content == frozen


def test_series_never_bypasses_legacy_final_approval(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "series-legacy", "admin", "*", "*")
        quotes = [fill(client, quote) for quote in series(client, "SERIES-LEGACY")]
        db_module = importlib.import_module("app.db")
        models = importlib.import_module("app.models.internal_quote")
        with db_module.SessionLocal() as db:
            legacy = db.get(models.InternalQuote, quotes[1]["id"])
            legacy.module_version = "v2"
            legacy.status = "ready_for_final_review"
            for section in db.query(models.InternalQuoteSection).filter_by(quote_id=legacy.id):
                if section.is_required:
                    section.status = "approved"
            db.commit()
        response = client.post(f"/api/internal-quotes/{quotes[0]['id']}/series-export", json=request_body(quotes))
        assert response.status_code == 409 and "最终审核" in response.text
        for quote in quotes:
            assert client.get(f"/api/internal-quotes/{quote['id']}/exports").json() == []
