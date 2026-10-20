"""Synthetic regression for Huakang A's per-order business version notes."""
from copy import deepcopy
from io import BytesIO
import json

import openpyxl
import pytest

from app.services import customer_order_unified as service
from app.services.huakang_a_order_legacy import unified_schedule as mapping
from app.schemas.customer_order import CustomerOrderLineOut
from test_customer_order_huakang_a_unified import template
from test_customer_order_unified import _parsed_row
from test_customer_order_ledger import engine, db, client, user


def workbook_bytes(workbook):
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def parsed_row():
    row = _parsed_row()
    row.update(id='360-test-1', reference_no='RL-765431-23', po_no='RL-765431-23',
               contract_no='RL765432122', product_no='7654002', customer_name='Example Customer',
               _huakang_customer='360', product_version='')
    return row


def kwargs(content):
    return dict(customer_code='360', factory_id='huakang-a', received_date='2026-10-15',
                po_files=[('po.xlsx', b'synthetic')], schedule_file_name='360.xlsx', schedule_content=content)


def parse_stub(monkeypatch):
    monkeypatch.setattr(service, '_parse_huakang_a_rows', lambda *_: ([parsed_row()], [], 'TEST'))


def test_version_can_be_entered_exported_on_all_sheets_and_serialized(monkeypatch):
    parse_stub(monkeypatch)
    content = template('360')
    preview = service.create_unified_customer_preview(**kwargs(content))
    row = preview['rows'][0]
    issue = next(i for i in row['issues'] if i['field'] == 'product_version')
    assert issue['severity'] == 'warning' and issue['can_edit']
    assert row['status'] != 'blocked'
    notes = '两个装\n20个样板；加拿大版本'
    overrides = [dict(row_id=row['id'], issue_key=issue['skip_key'], field='product_version', value=notes)]
    output, _, resolved = service.export_unified_customer_schedule(**kwargs(content),
        skipped_issue_keys={issue['skip_key']}, manual_overrides=overrides)
    assert resolved['rows'][0]['product_version'] == notes
    assert CustomerOrderLineOut.model_validate(resolved['rows'][0]).product_version == notes
    result = openpyxl.load_workbook(BytesIO(output))
    for name in service.SHEETS:
        sheet = result[name]
        columns = [c.column for c in sheet[3] if c.value == '版本']
        assert len(columns) == 1
        reference_column = 5 if name == service.ITEM_SHEET else 4
        output_row = next(c.row for c in sheet._cells.values() if c.column == reference_column and c.value == 'RL-765431-23')
        cell = sheet.cell(output_row, columns[0])
        assert cell.value == notes and cell.data_type == 's'
        assert cell.alignment.wrap_text and sheet.row_dimensions[output_row].height >= 30
    result.close()


def test_old_templates_append_without_shifting_existing_fields_or_history(monkeypatch):
    parse_stub(monkeypatch)
    workbook = openpyxl.load_workbook(BytesIO(template('360')))
    for name in service.SHEETS:
        sheet = workbook[name]
        for cell in sheet[3]:
            if cell.value == '版本':
                cell.value = '人工扩展列'
        sheet.cell(4, 45, '历史最末列不能移动')
        sheet.cell(4, 46, '=K4*2')
    content = workbook_bytes(workbook)
    output, _, _ = service.export_unified_customer_schedule(**kwargs(content))
    result = openpyxl.load_workbook(BytesIO(output))
    for name in service.SHEETS:
        sheet = result[name]
        assert sheet.cell(3, 47).value == '版本'
        assert sheet.cell(5, 47).value is None  # Empty is optional, not inferred.
        assert sheet.cell(4, 45).value == '历史最末列不能移动'
        assert sheet.cell(4, 46).value == '=K4*2'
        assert sheet.cell(4, 7).value is not None
    result.close()


def test_unique_existing_version_column_is_used_and_duplicate_labels_fail():
    workbook = openpyxl.load_workbook(BytesIO(template('360')))
    sheet = workbook[service.ITEM_SHEET]
    original = next(c for c in sheet[3] if c.value == '版本')
    original.value = '其他备注'
    sheet.cell(3, 42, '版 本')
    columns = mapping.ensure_360_version_columns(workbook)
    assert columns[service.ITEM_SHEET] == 42
    sheet.cell(3, 43, '版本')
    with pytest.raises(service.CustomerOrderUnifiedError, match='版本.*重复'):
        mapping.ensure_360_version_columns(workbook)
    workbook.close()


def test_version_history_is_exact_order_only_and_conflicts_remain_blank():
    history = [service.UnifiedHistoryRow(row=4, reference_no='RL-765431-23', product_no='7654002',
               extras={'product_version':'两个装\n20个样板'})]
    row = parsed_row()
    mapping.prepare_360_version(row, history)
    assert row['product_version'] == '两个装\n20个样板'
    fresh = parsed_row()
    fresh['reference_no'] = 'RL-765431-24'
    mapping.prepare_360_version(fresh, history)
    assert fresh['product_version'] == ''
    conflict = deepcopy(history[0])
    conflict.extras['product_version'] = '单个装'
    row = parsed_row()
    mapping.prepare_360_version(row, [*history, conflict])
    assert row['product_version'] == ''


def test_version_literal_not_formula_and_never_uses_po_revision():
    row = parsed_row()
    row['revision'] = 7
    mapping.prepare_360_version(row, [])
    assert row['product_version'] == ''
    row['product_version'] = '=HYPERLINK("https://example.invalid","版本")'
    workbook = openpyxl.load_workbook(BytesIO(template('360')))
    columns = mapping.ensure_360_version_columns(workbook)
    mapping.write_360_version(workbook[service.ITEM_SHEET], 5, columns[service.ITEM_SHEET], row)
    output = workbook_bytes(workbook)
    result = openpyxl.load_workbook(BytesIO(output))
    cell = result[service.ITEM_SHEET].cell(5, columns[service.ITEM_SHEET])
    assert cell.value == row['product_version'] and cell.data_type == 's'
    result.close()
    history = service.read_unified_history(output, factory_id='huakang-a', customer_code='360')
    # Give the written row the current order identity, then exercise real history.
    workbook = openpyxl.load_workbook(BytesIO(output))
    sheet = workbook[service.ITEM_SHEET]
    sheet.cell(5, 5, row['reference_no'])
    sheet.cell(5, 8, row['product_no'])
    output = workbook_bytes(workbook)
    history = service.read_unified_history(output, factory_id='huakang-a', customer_code='360')
    fresh = parsed_row()
    mapping.prepare_360_version(fresh, history)
    assert fresh['product_version'] == row['product_version']
    workbook = openpyxl.load_workbook(BytesIO(output))
    workbook[service.ITEM_SHEET].cell(5, columns[service.ITEM_SHEET]).data_type = 'f'
    history = service.read_unified_history(workbook_bytes(workbook), factory_id='huakang-a', customer_code='360')
    fresh = parsed_row()
    mapping.prepare_360_version(fresh, history)
    assert fresh['product_version'] == ''


@pytest.mark.parametrize('customer', ['green-toys', 'headstart'])
def test_other_customer_extensions_do_not_gain_version_column(customer):
    workbook = openpyxl.load_workbook(BytesIO(template(customer)))
    assert all(c.value != '版本' for c in workbook[service.ITEM_SHEET][3])
    assert 'product_version' not in mapping.EXTRA_HEADERS[customer]
    workbook.close()


def test_real_converted_excel_parser_and_http_import_keep_version_and_audit_controls(client, db):
    from app.services.auth import get_current_user
    from app.api import customer_order as api
    client.test_app.dependency_overrides[get_current_user] = lambda: user(['read', 'write'], factory='huakang-a')
    po = openpyxl.Workbook()
    for n, value in enumerate([
        'PURCHASE ORDER RELEASE (FIRM) RL-765431-23', 'REVISION 7',
        'Revision Date: 15-Oct-2026', 'Customer PO Number: 7654321',
        'Customer Release No.: RL765432122', 'End Customer: Example Customer',
        '1 7654002 Each 360', 'Toy Example Robot', 'Master Carton Qty: 2',
        'Planned Inspection Date: 15-Oct-2026', 'Factory Commit Date: 22-Oct-2026',
    ], 1):
        po.active.cell(n, 1, value)
    args = kwargs(template('360'))
    args['po_files'] = [('po.xlsx', workbook_bytes(po))]
    raw = service.create_unified_customer_preview(**args)
    api._finalize_preview(raw, received_date=args['received_date'], current_user=user(['read', 'write'], factory='huakang-a'), factory_id='huakang-a')
    row = raw['rows'][0]
    assert row['product_no'] == '7654002' and row['product_version'] == ''
    issue = next(i for i in row['issues'] if i['field'] == 'product_version')
    notes = '两个装\n业务补充20个样板'
    overrides = [dict(row_id=row['id'], issue_key=issue['skip_key'], field='product_version', value=notes)]
    form = dict(factory_id='huakang-a', received_date=args['received_date'], confirmed='true',
                preview_fingerprint=raw['preview_fingerprint'], manual_overrides=json.dumps(overrides),
                confirmation_reason='按业务原件核对版本备注')
    files = [('po_files', ('po.xlsx', args['po_files'][0][1], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')),
             ('schedule_file', ('360.xlsx', args['schedule_content'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))]
    response = client.post('/api/customer-order-ledger/imports/360', data=form, files=files)
    assert response.status_code == 200, response.text
    saved = response.json()['items'][0]
    assert saved['data']['product_version'] == notes
    assert saved['data']['import_controls']['manual_overrides'] == overrides
    assert saved['data']['import_controls']['reason'] == form['confirmation_reason']
    assert saved['version'] == 1
    form['confirmation_reason'] = ''
    assert client.post('/api/customer-order-ledger/imports/360', data=form, files=files).status_code == 400
