from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest

from app.services import customer_order_unified as service
from app.services.huakang_a_order_legacy import green_toys_headstart as parser
from app.services.huakang_a_order_legacy import unified_schedule as mapping


LINES = [
    ('DTK01R', 'Dump Truck Red', '2.847', '14,235.00'),
    ('DTKBO-1283', 'Dump Truck - Blue and Orange', '2.847', '14,235.00'),
    ('DTKP-1010', 'Dump Truck - Pink', '2.847', '14,235.00'),
    ('HELG-1061', 'Helicopter - Green', '2.015', '10,075.00'),
    ('SUBY-1033', 'Submarine - Yellow', '1.315', '6,575.00'),
    ('SUBB-1032', 'Submarine - Blue', '1.315', '6,575.00'),
    ('WAGO-1227', 'Wagon - Orange', '3.533', '17,665.00'),
    ('HELB-1060', 'Helicopter - Blue', '2.015', '10,075.00'),
]
HEADER = 'PO Date 10/1/2026\nPO # 7848\nDeliver By Date 11/10/2026\n'


def html_po():
    rows = ''.join('<tr>' + ''.join(f'<td>{v}</td>' for v in (item, name, '5,000', '0', '', price, amount)) + '</tr>'
                   for item, name, price, amount in LINES)
    return ('<html><body>Green Toys, Inc<script>throw new Error("must not execute")</script>'
            '<table><tr><td>PO Date</td><td></td><td>10/1/2026</td></tr>'
            '<tr><td>PO #</td><td></td><td>7848</td></tr>'
            '<tr><td>Deliver By Date</td><td></td><td>11/10/2026</td></tr></table>'
            '<table><tr><th>Item</th><th>Description</th><th>Order Quantity</th><th>Received</th>'
            '<th>Inventory Detail</th><th>Unit Price</th><th>Amount</th></tr>' + rows +
            '<tr><td colspan="6">Total</td><td>$93,670.00</td></tr></table></body></html>').encode()


def test_native_html_preserves_all_eight_rows_and_totals():
    rows = parser.parse_green_toys_po('Purchase Order_7848.html', html_po())
    assert [r['item_no'] for r in rows] == [r[0] for r in LINES]
    assert all(r['parse_ok'] and not r['parse_warnings'] for r in rows)
    assert sum(r['quantity'] for r in rows) == 40000
    assert sum(Decimal(str(r['amount_usd'])) for r in rows) == Decimal('93670')
    assert rows[0]['order_date'] == '2026-10-01'
    assert rows[0]['factory_commit_date'] == '2026-11-10'
    assert rows[0]['planned_inspection_date'] == '2026-11-09'


@pytest.mark.parametrize('source', [
    html_po().replace(b'17,665.00', b'17,664.00'),
    html_po().replace(b'$93,670.00', b'$90,000.00'),
    html_po().replace(b'<td>SUBY-1033</td>', b'<td></td>'),
    html_po().replace(b'<td>5,000</td>', b'<td>5O00</td>', 1),
    html_po().replace(b'<td>11/10/2026</td>', b'<td>invalid</td>'),
], ids=['line-amount', 'total-amount', 'missing-item', 'invalid-number', 'invalid-date'])
def test_native_html_rejects_incomplete_or_conflicting_evidence(source):
    with pytest.raises(parser.HuakangASpecialCustomerError):
        parser.parse_green_toys_po('7848.html', source)


def test_single_space_ocr_does_not_drop_three_lines():
    lines = [(' ' if i in {4, 5, 6} else '     ').join((item, desc, '5,000', '0', price, amount))
             for i, (item, desc, price, amount) in enumerate(LINES)]
    rows = parser.parse_green_toys_ocr_text('7848.png', HEADER + '\n'.join(lines) + '\nTotal $93,670.00')
    assert [r['item_no'] for r in rows] == [r[0] for r in LINES]
    assert all(r['parse_ok'] for r in rows)


def test_ocr_unreadable_candidate_remains_visible_and_blocked():
    text = HEADER + 'SUBY-1033 Submarine - Yellow unreadable quantity and price\n'
    rows = parser.parse_green_toys_ocr_text('7848.png', text)
    assert len(rows) == 1 and rows[0]['item_no'] == 'SUBY-1033'
    assert not rows[0]['parse_ok']
    issues = parser.issues_for_record('green-toys', 'Green Toys', rows[0], '1')
    assert any(i['severity'] == 'blocked' for i in issues)
    assert len({i['skip_key'] for i in issues}) == len(issues)


def test_ocr_total_mismatch_rejects_partial_batch():
    text = HEADER + 'DTK01R     Dump Truck Red     5,000     0     2.847     14,235.00\nTotal $93,670.00'
    with pytest.raises(parser.HuakangASpecialCustomerError, match='总金额'):
        parser.parse_green_toys_ocr_text('7848.png', text)


def test_ocr_lost_decimal_uses_amount_only_when_scale_is_supported():
    text = HEADER + 'WAGO-1227 Wagon - Orange 5,000 0 3533 17,665.00\nTotal $17,665.00'
    rows = parser.parse_green_toys_ocr_text('7848.png', text)
    assert rows[0]['unit_price_usd'] == 3.533
    assert rows[0]['parse_ok']


def recovery_template():
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, marker in ((service.ITEM_SHEET, 8), (service.ORDER_SHEET, 7), (service.REVIEW_SHEET, 9)):
        sheet = book.create_sheet(name)
        headers = service.ITEM_HEADERS[:29] + mapping.ITEM_TAILS['green-toys'] if name == service.ITEM_SHEET else service.ORDER_HEADERS
        for c, value in enumerate(headers, 1):
            sheet.cell(3, c, value)
        sheet.cell(marker, 1, '取消单')
        sheet.cell(marker + 1, 1, '保留旧内容')
        if name == service.ITEM_SHEET:
            sheet['F4'], sheet['H4'], sheet['K4'] = '7816-14-1', 'OLD', 2500
            sheet['AH4'] = 1.23
            for n, (item, desc, _, _) in enumerate(LINES, 20):
                sheet.cell(n, 6, f'7000-{n}')
                sheet.cell(n, 8, item)
                sheet.cell(n, 9, '历史中文品名')
                sheet.cell(n, 10, desc)
                sheet.cell(n, 11, 10)
                sheet.cell(n, 12, 4 if item == 'WAGO-1227' else 8)
        else:
            for n, ref, item in ((4, '7848-1', 'DTK01R'), (5, '7848-5', 'HELB-1060')):
                for c, v in enumerate(('7848', ref, ref, 'Green Toys', item), 3):
                    sheet.cell(n, c, v)
                for c, ic in ((1, 1), (8, 9), (10, 11), (18, 25)):
                    sheet.cell(n, c, f'=ITEM表!{openpyxl.utils.get_column_letter(ic)}{n}')
                sheet.cell(n, 14, 99)  # user-owned factory price
                sheet.cell(n, 15, f'=J{n}*N{n}')
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def preview(content):
    return service.create_unified_customer_preview(customer_code='green-toys', factory_id='huakang-a',
        received_date='2026-10-05', po_files=[('7848.html', html_po())],
        schedule_file_name='Green Toys.xlsx', schedule_content=content)


def test_recovery_preserves_five_suffix_and_manual_price_without_summary_duplicates():
    content = recovery_template()
    result = preview(content)
    refs = {r['product_no']: r['po_no'] for r in result['rows']}
    assert refs['DTK01R'] == '7848-1'
    assert refs['HELB-1060'] == '7848-5'
    assert refs['SUBY-1033'] == '7848-9'
    assert result['summary']['total'] == 8 and result['summary']['blocked'] == 0
    output, _, _ = service.export_unified_customer_schedule(customer_code='green-toys', factory_id='huakang-a',
        received_date='2026-10-05', po_files=[('7848.html', html_po())],
        schedule_file_name='Green Toys.xlsx', schedule_content=content)
    book = openpyxl.load_workbook(BytesIO(output))
    item = book[service.ITEM_SHEET]
    details = {row[7].value: row[0].row for row in item if row[3].value == '7848'}
    assert len(details) == 8
    assert item['F4'].value == '7816-14-1' and item['K4'].value == 2500
    for name in (service.ORDER_SHEET, service.REVIEW_SHEET):
        sheet = book[name]
        summaries = [row for row in sheet if row[2].value == '7848']
        assert len(summaries) == 8
        assert sheet['N4'].value == sheet['N5'].value == 99
        for row in summaries:
            detail = details[row[6].value]
            assert row[3].value == item.cell(detail, 5).value
            assert row[9].value == f'=IF(LEN(\'ITEM表\'!K{detail})=0,"",\'ITEM表\'!K{detail})'
        assert any(cell.value == '保留旧内容' for cell in sheet['A'])
    # A new import sees restored ITEM history and retains controlled duplicates.
    repeated = preview(output)
    assert repeated['summary']['blocked'] == 8
    assert all(any(i['code'] == 'duplicate_existing_order' for i in r['issues']) for r in repeated['rows'])


@pytest.mark.parametrize('change', ['missing-peer', 'literal-quantity', 'different-product', 'duplicate-reference', 'reference-in-use'])
def test_recovery_refuses_conflicting_or_manual_summary_fields(change):
    book = openpyxl.load_workbook(BytesIO(recovery_template()))
    if change == 'missing-peer':
        book[service.REVIEW_SHEET]['C4'] = None
    elif change == 'literal-quantity':
        book[service.ORDER_SHEET]['J4'] = 6000
    elif change == 'different-product':
        book[service.REVIEW_SHEET]['G4'] = 'HELG-1061'
    elif change == 'duplicate-reference':
        for name in (service.ORDER_SHEET, service.REVIEW_SHEET):
            book[name]['D5'] = book[name]['E5'] = '7848-1'
    else:
        item = book[service.ITEM_SHEET]
        item['F4'], item['H4'] = '7848-1', 'SUBY-1033'
    buffer = BytesIO()
    book.save(buffer)
    with pytest.raises(service.CustomerOrderUnifiedError):
        preview(buffer.getvalue())


def test_native_html_keeps_factory_isolation():
    with pytest.raises(service.CustomerOrderUnifiedError, match='不属于'):
        service.create_unified_customer_preview(customer_code='green-toys', factory_id='huaxing',
            received_date='2026-10-05', po_files=[('7848.html', html_po())],
            schedule_file_name='Green Toys.xlsx', schedule_content=recovery_template())


def test_export_guard_refuses_missing_item_write(monkeypatch):
    monkeypatch.setattr(service, '_write_item_row', lambda *_args: None)
    with pytest.raises(service.CustomerOrderUnifiedError, match='三表身份不一致'):
        service.export_unified_customer_schedule(customer_code='green-toys', factory_id='huakang-a',
            received_date='2026-10-05', po_files=[('7848.html', html_po())],
            schedule_file_name='Green Toys.xlsx', schedule_content=recovery_template())
