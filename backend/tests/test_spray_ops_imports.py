"""Source fidelity and document-boundary replay protection using synthetic files."""
from io import BytesIO
import json
from uuid import uuid4

import openpyxl
import pytest
import xlwt

from test_spray_ops import client, payload, post, read, user_with
from app.services.spray_ops.source_reader import read_source, PROFILES


def workbook(rows=None, *, marker=""):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "订单"
    for row in rows or [["单号", "货号", "数量", "单价"], ["TEST-001", "00017", 100, "0.26"]]:
        sheet.append(row)
    if marker:
        sheet["G1"] = marker
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def upload(client, content, profile="spray-intake-v1", filename="deliberately.et"):
    response = client.post("/api/spray-operations/imports/upload", data={"options": json.dumps(payload(profile=profile))}, files={"file": (filename, content, "application/octet-stream")})
    assert response.status_code == 200, response.text
    return response.json()["data"]


def document(key="TEST-001"):
    return dict(document_key=key, sheet="订单", coordinates=["A2", "B2", "C2", "D2"], target="demand", values=dict(document_no=key, counterparty="合成往来方", business_date="2026-09-02", lines=[dict(item_no="00017", part="左", color="白", quantity="100", due_date="2026-09-30")]), reason="逐单人工核对", primary_source_confirmation="此工作簿为该订单主事实来源")


def test_signature_sparse_zeros_formula_and_date_preservation():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet["A1"] = 17
    sheet["A1"].number_format = "00000"
    sheet["B2"] = "=1/0"
    sheet["C3"] = "#REF!"
    sheet["XFD1000"].number_format = "0"
    output = BytesIO()
    book.save(output)
    parsed = read_source(output.getvalue())
    cells = {cell["coordinate"]: cell for cell in parsed["sheets"][0]["cells"]}
    assert parsed["format"] == "OOXML" and len(cells) == 3
    assert cells["A1"]["formatted"] == "00017"
    assert cells["B2"]["formula"] == "=1/0" and cells["B2"]["cached"] is None
    assert "formula_cache_missing" in cells["B2"]["flags"]
    assert cells["C3"]["error"] == "#REF!"
    legacy = xlwt.Workbook()
    legacy.add_sheet("入库").write(0, 0, "00017")
    output = BytesIO()
    legacy.save(output)
    assert read_source(output.getvalue())["format"] == "BIFF8"
    assert len(PROFILES) == 12


def test_import_per_document_confirm_retry_and_semantic_duplicates(client):
    source = upload(client, workbook())
    preview = post(client, "imports/preview", source_id=source["id"], profile="spray-intake-v1", business_month="2026-09", documents=[document()])
    body = payload(expected_version=preview["version"], fingerprint=preview["fingerprint"], document_key="TEST-001")
    url = f"/api/spray-operations/imports/{preview['id']}/confirm"
    result = client.post(url, json=body)
    assert result.status_code == 200, result.text
    assert client.post(url, json=body).json()["data"] == result.json()["data"]
    assert len(read(client, "demands")) == 1
    another = upload(client, workbook(marker="different physical workbook"))
    duplicate = post(client, "imports/preview", source_id=another["id"], profile="spray-intake-v1", business_month="2026-09", documents=[document()])
    assert duplicate["snapshot"]["documents"][0]["conflict_id"]
    failed = client.post(f"/api/spray-operations/imports/{duplicate['id']}/confirm", json=payload(expected_version=duplicate["version"], fingerprint=duplicate["fingerprint"], document_key="TEST-001"))
    assert failed.status_code == 409 and failed.json()["code"] == "duplicate_source_fact"


def test_import_requires_target_permission_and_protects_payroll_source(client):
    source = upload(client, workbook(), "spray-daily-payroll-v1")
    client.spray_auth["user"] = user_with("read", "import", "cost_read")
    assert client.get(f"/api/spray-operations/sources/{source['id']}", params={"factory_id": "huaxing"}).status_code == 403
    assert client.get(f"/api/spray-operations/sources/{source['id']}/download", params={"factory_id": "huaxing"}).status_code == 403
    order_source = upload(client, workbook(marker="nonpayroll"))
    preview = post(client, "imports/preview", source_id=order_source["id"], profile="spray-intake-v1", business_month="2026-09", documents=[document()])
    result = client.post(f"/api/spray-operations/imports/{preview['id']}/confirm", json=payload(expected_version=1, fingerprint=preview["fingerprint"], document_key="TEST-001"))
    assert result.status_code == 403


def test_negative_source_cannot_be_normal_production(client):
    source = upload(client, workbook([["单号", "货号", "数量", "单价"], ["TEST-001", "00017", -1800, "1.51"]]))
    result = client.post("/api/spray-operations/imports/preview", json=payload(source_id=source["id"], profile="spray-intake-v1", business_month="2026-09", documents=[document()]))
    assert result.status_code == 422 and result.json()["code"] == "negative_source"
