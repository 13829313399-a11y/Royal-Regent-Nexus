"""Dickie-only insertion and conditional-format regression, no business files."""
from copy import copy
from io import BytesIO

import openpyxl
import pytest
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.services import customer_order_unified as unified
from app.services import customer_order_huaxing_unified as mapping

SHEETS = ('ITEM表', mapping.DICKIE_EXTRA, '接单表', '正单评审表')


def columns(sheet):
    return (4, 7, 10) if sheet.title in ('接单表', '正单评审表') else (5, 8, 11)


def fixture_workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for index, name in enumerate(SHEETS):
        sheet = wb.create_sheet(name)
        summary = name in ('接单表', '正单评审表')
        for col, label in enumerate(unified.ORDER_HEADERS if summary else unified.ITEM_HEADERS, 1):
            sheet.cell(3, col, label)
        ref, product, quantity = columns(sheet)
        offset = 1 if summary else 0
        for r, item, identity, kind in ((4+offset,'20 374 7020','SC-OLD-A','正单'),
                (5+offset,'20 374 7020','已入系统','MO单'),
                (6+offset,'20 374 7020','已入系统','MO单'),
                (8+offset,'20 375 7005','SC-OLD-B','正单'),
                (9+offset,'20 375 7005','已入系统','MO单')):
            sheet.cell(r, 2, kind)
            sheet.cell(r, ref, identity)
            sheet.cell(r, product, item)
            sheet.cell(r, quantity, 120)
            sheet.cell(r, 40, '手工历史备注-'+identity)
            sheet.cell(r, 41, f'={openpyxl.utils.get_column_letter(quantity)}{r}*2')
            sheet.row_dimensions[r].height = 24 if kind=='正单' else 30
            for cell in sheet[r]:
                cell.font = Font(name='Arial', size=11, color='000000')
                if kind=='MO单':
                    cell.fill = PatternFill('solid', fgColor='A9D18E')
        marker=15+index
        sheet.cell(marker, 1, '取消单')
        sheet.cell(marker+1, ref, 'SC-CANCELLED')
        sheet.cell(marker+1, product, '20 374 7020')
        sheet.cell(marker+1, quantity, 999)
        sheet.cell(marker+3, 1, '已走货订单')
        sheet.merge_cells(start_row=marker,start_column=1,end_row=marker,end_column=3)
        sheet.freeze_panes='H4'
        validation=DataValidation(type='list',formula1='"正单,MO单"')
        sheet.add_data_validation(validation)
        validation.add(f'B4:B{marker-1}')
        # Real source has this E/H rule on all sheets, including the summaries.
        sheet.conditional_formatting.add(f'E4:H{marker+1}', FormulaRule(
            formula=[f'AND($E4<>"",$H4<>"",COUNTIFS($E$4:$E${marker+1},$E4,$H$4:$H${marker+1},$H4)>1)'],
            font=Font(color='FFB4232E')))
    return wb


def record(ref, product, *, lane='ITEM表', duplicate=False):
    return dict(id=ref, reference_no=ref, po_no=ref+'-PO', contract_no=ref+'-CT',
                product_no=product, quantity='24', units_per_carton='6',
                customer_name='Simba', product_name_zh='小邮递车', product_name_en='Toy',
                order_type='正单', item_sheet_name=lane, _huaxing_customer='dickie',
                issues=[{'code':'duplicate_order_warning'}] if duplicate else [])


def find(sheet, ref):
    ref_col, _, _=columns(sheet)
    return next(r for r in range(4,sheet.max_row+1) if sheet.cell(r,ref_col).value==ref)


def test_new_details_group_by_complete_item_above_stock_and_link_actual_rows():
    wb=fixture_workbook()
    rows=[record('SC-NEW-B1','203757005'), record('SC-NEW-A','203747020'),
          record('SC-NEW-B2','20 375 7005')]
    mapping.export_rows(wb,rows)
    for name in ('ITEM表','接单表','正单评审表'):
        sheet=wb[name]
        ref_col, product_col, _=columns(sheet)
        old_a=find(sheet,'SC-OLD-A')
        new_a=find(sheet,'SC-NEW-A')
        new_b1=find(sheet,'SC-NEW-B1')
        new_b2=find(sheet,'SC-NEW-B2')
        assert new_a==old_a+1
        assert new_b2==new_b1+1
        assert new_b1==find(sheet,'SC-OLD-B')+1
        assert sheet.cell(new_a+1,2).value=='MO单'
        assert sheet.cell(new_a,ref_col).font.color.rgb=='FF0000FF'
        assert sheet.cell(new_a,product_col).fill.patternType is None
        assert sheet.cell(old_a,40).value=='手工历史备注-SC-OLD-A'
        assert sheet.cell(old_a,41).value==f'={openpyxl.utils.get_column_letter(columns(sheet)[2])}{old_a}*2'
        assert sheet.freeze_panes=='H4'
        assert find(sheet,'SC-CANCELLED')>unified._marker_row(sheet)
    for reference in ('SC-NEW-A','SC-NEW-B1','SC-NEW-B2'):
        item_row=find(wb['ITEM表'],reference)
        for name in ('接单表','正单评审表'):
            r=find(wb[name],reference)
            assert wb[name].cell(r,10).value==f'=IF(LEN(\'ITEM表\'!K{item_row})=0,"",\'ITEM表\'!K{item_row})'


def test_stock_only_group_uses_normal_style_and_other_item_suffix_is_not_merged():
    wb=fixture_workbook()
    for name in ('ITEM表','接单表','正单评审表'):
        s=wb[name]
        s.cell(find(s,'SC-OLD-A'),columns(s)[1],'DIFFERENT')
    mapping.export_rows(wb,[record('SC-NEW-A','203747020'),record('SC-NEW-C','203747028')])
    for name in ('ITEM表','接单表','正单评审表'):
        s=wb[name]
        assert find(s,'SC-NEW-A')<find(s,'SC-OLD-B')
        assert find(s,'SC-NEW-C')>find(s,'SC-OLD-B')
        assert s.cell(find(s,'SC-NEW-A'),columns(s)[1]).fill.patternType is None


def test_only_business_confirmed_series_members_share_group_but_keep_full_codes():
    wb=fixture_workbook()
    mapping.export_rows(wb,[record('SC-SERIES','20 374 7029')])
    for name in ('ITEM表','接单表','正单评审表'):
        s=wb[name]
        r=find(s,'SC-SERIES')
        assert r==find(s,'SC-OLD-A')+1
        assert s.cell(r,columns(s)[1]).value=='20 374 7029'
        assert s.cell(r+1,2).value=='MO单'


def test_dinosaur_lane_is_independent_and_duplicate_font_is_retained():
    wb=fixture_workbook()
    mapping.export_rows(wb,[record('SC-DINO','203747020',lane=mapping.DICKIE_EXTRA,duplicate=True)])
    for name in (mapping.DICKIE_EXTRA,'接单表','正单评审表'):
        s=wb[name]
        r=find(s,'SC-DINO')
        assert r==find(s,'SC-OLD-A')+1
        assert s.cell(r,columns(s)[0]).font.color.rgb=='FFFF0000'
    assert all(c.value!='SC-DINO' for c in wb['ITEM表']['E'])


def test_duplicate_color_rules_exclude_status_text_and_mo_rows_and_are_idempotent():
    wb=fixture_workbook()
    mapping._repair_dickie_duplicate_rules(wb)
    before={}
    for name in SHEETS:
        s=wb[name]
        rule=next(rule for rules in s.conditional_formatting._cf_rules.values() for rule in rules if rule.formula)
        formula=rule.formula[0]
        ref,item,_=columns(s)
        ref_letter,item_letter=map(openpyxl.utils.get_column_letter,(ref,item))
        assert f'LEFT(UPPER(TRIM(${ref_letter}4)),2)="SC"' in formula
        assert f'${item_letter}4<>""' in formula
        assert 'SEARCH("MO",$B4)' in formula
        assert 'SEARCH("备料",$B4)' in formula
        assert f'${ref_letter}$4:' in formula and f'${item_letter}$4:' in formula
        before[name]=formula
    mapping._repair_dickie_duplicate_rules(wb)
    assert before=={name:next(rule.formula[0] for rules in wb[name].conditional_formatting._cf_rules.values() for rule in rules if rule.formula) for name in SHEETS}


def test_unrelated_conditional_rules_and_historical_cell_styles_remain():
    wb=fixture_workbook()
    s=wb['ITEM表']
    s.conditional_formatting.add('K4:K14', FormulaRule(formula=['K4<0'],font=Font(color='ABCDEF')))
    old=copy(s['H5']._style)
    mapping.export_rows(wb,[record('SC-NEW-A','203747020')])
    stock=find(s,'已入系统')
    assert s.cell(stock,8)._style==old
    assert any(rule.formula==['K4<0'] for rules in s.conditional_formatting._cf_rules.values() for rule in rules)


@pytest.mark.parametrize('formula',[
    'AND($D4<>"",$G4<>"",COUNTIFS($D$4:$D$20,$D4,$G$4:$G$20,$G4)>1)',
    'AND($D4<>"",$H4<>"",COUNTIFS($D$4:$D$20,$D4,$H$4:$H$20,$H4)>1)',
    'AND($E4<>"",$H4<>"",COUNTIFS($D$4:$D$20,$E4,$H$4:$H$20,$H4)>1)',
])
def test_other_item_lane_key_pairs_are_not_misidentified_as_so_item_rules(formula):
    wb=fixture_workbook()
    for name in ('ITEM表',mapping.DICKIE_EXTRA):
        wb[name].conditional_formatting.add('D4:H10',FormulaRule(formula=[formula],font=Font(color='ABCDEF')))
    mapping._repair_dickie_duplicate_rules(wb)
    for name in ('ITEM表',mapping.DICKIE_EXTRA):
        assert any(rule.formula==[formula] for rules in wb[name].conditional_formatting._cf_rules.values() for rule in rules)


@pytest.mark.parametrize('old,new',[('>1)', '>5)'), (',2)="SC"', ',3)="SC"')])
def test_lookalike_custom_rule_constants_are_not_rewritten(old,new):
    wb=fixture_workbook()
    mapping._repair_dickie_duplicate_rules(wb)
    formulas={}
    for name in SHEETS:
        rule=next(rule for rules in wb[name].conditional_formatting._cf_rules.values() for rule in rules if rule.formula)
        rule.formula=[rule.formula[0].replace(old,new)]
        formulas[name]=rule.formula[:]
    mapping._repair_dickie_duplicate_rules(wb)
    assert formulas=={name:next(rule.formula for rules in wb[name].conditional_formatting._cf_rules.values() for rule in rules if rule.formula) for name in SHEETS}


def test_saved_workbook_roundtrip_keeps_grouping_and_number_of_rows():
    wb=fixture_workbook()
    mapping.export_rows(wb,[record('SC-NEW-A','203747020')])
    out=BytesIO(); wb.save(out); wb.close()
    saved=openpyxl.load_workbook(BytesIO(out.getvalue()))
    assert find(saved['ITEM表'],'SC-NEW-A')==5
    assert sum(c.value=='SC-NEW-A' for c in saved['ITEM表']['E'])==1
    assert saved['ITEM表'].cell(find(saved['ITEM表'],'SC-CANCELLED'),11).value==999
    saved.close()


def test_new_order_does_not_join_matching_cancelled_or_shipped_history():
    wb=fixture_workbook()
    for name in SHEETS:
        s=wb[name]
        ref,product,quantity=columns(s)
        boundary=unified._marker_row(s)
        for r in range(4,boundary):
            if s.cell(r,product).value=='20 374 7020':
                s.cell(r,product,'OTHER-ACTIVE-PRODUCT')
        s.cell(boundary+4,ref,'SC-SHIPPED')
        s.cell(boundary+4,product,'20 374 7020')
        s.cell(boundary+4,quantity,456)
    mapping.export_rows(wb,[record('SC-NEW-A','203747020')])
    for name in ('ITEM表','接单表','正单评审表'):
        s=wb[name]
        marker=unified._marker_row(s)
        assert find(s,'SC-NEW-A')<marker<find(s,'SC-CANCELLED')<find(s,'SC-SHIPPED')
        assert s.cell(find(s,'SC-CANCELLED'),columns(s)[2]).value==999
        assert s.cell(find(s,'SC-SHIPPED'),columns(s)[2]).value==456


@pytest.mark.parametrize('title',['已走货订单','已完成订单'])
def test_first_historical_section_stops_active_grouping_even_before_cancellation(title):
    wb=fixture_workbook()
    for name in ('ITEM表','接单表','正单评审表'):
        s=wb[name]
        s.cell(find(s,'SC-OLD-B')-1,1,title)
    mapping.export_rows(wb,[record('SC-NEW-B','203757005')])
    for name in ('ITEM表','接单表','正单评审表'):
        s=wb[name]
        banner=next(c.row for c in s['A'] if c.value==title)
        assert find(s,'SC-NEW-B')<banner<find(s,'SC-OLD-B')


def test_order_notes_containing_cancellation_words_are_not_section_titles():
    wb=fixture_workbook()
    for name in ('ITEM表','接单表','正单评审表'):
        s=wb[name]
        s.cell(4,40,'客户要求核对取消单后再安排')
        # A source value at a left-hand data field is not a section banner.
        s.cell(find(s,'SC-OLD-A'),1,'取消单')
    mapping.export_rows(wb,[record('SC-NEW-B','203757005')])
    for name in ('ITEM表','接单表','正单评审表'):
        s=wb[name]
        assert find(s,'SC-NEW-B')==find(s,'SC-OLD-B')+1
        assert s.cell(4,40).value=='客户要求核对取消单后再安排'


def test_unknown_product_summary_is_inserted_before_annual_totals():
    wb=fixture_workbook()
    for name in ('接单表','正单评审表'):
        s=wb[name]
        s.cell(12,10,'2026年接单')
        s.cell(13,10,'2025-2026年接单合计')
    mapping.export_rows(wb,[record('SC-NEW-UNKNOWN','203747028')])
    for name in ('接单表','正单评审表'):
        s=wb[name]
        total=next(c.row for c in s['J'] if c.value=='2026年接单')
        assert find(s,'SC-NEW-UNKNOWN')<total


def test_blank_display_formula_keys_do_not_hide_historical_section_title():
    wb=fixture_workbook()
    for name in ('ITEM表','接单表','正单评审表'):
        s=wb[name]
        banner=find(s,'SC-OLD-B')-1
        s.cell(banner,1,'已走货订单')
        for column in columns(s)[:2]:
            s.cell(banner,column,'=""')
    mapping.export_rows(wb,[record('SC-NEW-B','203757005')])
    for name in ('ITEM表','接单表','正单评审表'):
        s=wb[name]
        banner=next(c.row for c in s['A'] if c.value=='已走货订单')
        assert find(s,'SC-NEW-B')<banner<find(s,'SC-OLD-B')
        assert all(s.cell(banner,c).value=='=""' for c in columns(s)[:2])
