"""Synthetic regressions for Shushupapa's first, active order section."""
from copy import copy
from io import BytesIO

import openpyxl
import pytest
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

from app.services import customer_order_huaxing_unified as huaxing
from app.services import customer_order_unified as unified
from app.services.huakang_a_order_legacy import unified_schedule as slots
from test_customer_order_shushupapa_export import find_row, order


NAMES = ("ITEM表", "接单表", "正单评审表")


def book_with_sections(totals=(5, 5, 5), *, subtotal=True):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, total in zip(NAMES, totals):
        sheet = book.create_sheet(name)
        headers = unified.ITEM_HEADERS if name == "ITEM表" else unified.ORDER_HEADERS
        for column, header in enumerate(headers, 1):
            sheet.cell(3, column, header)
        quantity = "K" if name == "ITEM表" else "J"
        if subtotal:
            sheet.cell(total, 1, "合计：")
            sheet[f"{quantity}{total}"] = f"=SUM({quantity}4:{quantity}{total - 1})"
            sheet.cell(total + 1, 1, "备料单")
        else:
            sheet.cell(total, 1, "备料单")
        # Reservations and historical orders are deliberately below the first
        # section, but above cancellation. They must not influence new placement.
        sheet.cell(total + 3, 4, f"SYNTH-RESERVATION-{name}")
        sheet.cell(total + 3, 8, "SYNTH-RESERVED-ITEM")
        sheet.cell(total + 3, 11, 400)
        sheet.cell(total + 5, 1, "取消单")
        sheet.cell(total + 6, 4, f"SYNTH-CANCELLED-{name}")
        sheet.cell(total + 8, 1, "已走货订单")
        sheet.cell(total + 9, 4, f"SYNTH-SHIPPED-{name}")
    return book


def row_with_value(sheet, value):
    return next(cell.row for cell in sheet._cells.values() if cell.value == value)


def reopened(book):
    data = BytesIO()
    book.save(data)
    return openpyxl.load_workbook(BytesIO(data.getvalue()))


def historical_order(sheet, number, reference):
    if sheet.title == "ITEM表":
        sheet.cell(number, 4, reference)
        sheet.cell(number, 5, reference)
        sheet.cell(number, 8, f"ITEM-{reference}")
    else:
        sheet.cell(number, 3, reference)
        sheet.cell(number, 4, reference)
        sheet.cell(number, 7, f"ITEM-{reference}")


def test_six_orders_expand_the_single_top_row_above_its_subtotal():
    book = book_with_sections()
    before = {
        name: {r: tuple(book[name].cell(r, c).value for c in range(1, 12))
               for r in range(6, 15)}
        for name in NAMES
    }
    rows = [order(i) for i in range(1, 7)]
    huaxing.export_rows(book, rows)
    book = reopened(book)
    for name in NAMES:
        sheet = book[name]
        assert [find_row(sheet, row["reference_no"]) for row in rows] == list(range(4, 10))
        assert sheet["A10"].value == "合计："
        quantity = "K" if name == "ITEM表" else "J"
        assert sheet[f"{quantity}10"].value == f"=SUM({quantity}4:{quantity}9)"
        for number, values in before[name].items():
            assert tuple(sheet.cell(number + 5, c).value for c in range(1, 12)) == values


def test_each_sheet_uses_its_own_top_area_and_summary_formulas_follow_actual_item_rows():
    book = book_with_sections((5, 7, 6))
    historical_order(book["接单表"], 4, "SYNTH-OLD-ORDER")
    historical_order(book["正单评审表"], 4, "SYNTH-OLD-REVIEW-1")
    historical_order(book["正单评审表"], 5, "SYNTH-OLD-REVIEW-2")
    rows = [order(i) for i in range(1, 7)]
    huaxing.export_rows(book, rows)
    expected = {"ITEM表": range(4, 10), "接单表": range(5, 11), "正单评审表": range(6, 12)}
    for name, numbers in expected.items():
        sheet = book[name]
        assert [find_row(sheet, row["reference_no"]) for row in rows] == list(numbers)
        assert sheet.cell(numbers.stop, 1).value == "合计："
    for index, row in enumerate(rows):
        item_number, order_number, review_number = 4 + index, 5 + index, 6 + index
        ref = f"'ITEM表'!K{item_number}"
        for name, number in (("接单表", order_number), ("正单评审表", review_number)):
            assert book[name].cell(number, 10).value == f'=IF(LEN({ref})=0,"",{ref})'
        price_ref = f"'接单表'!P{order_number}"
        assert book["正单评审表"].cell(review_number, 16).value == f'=IF(LEN({price_ref})=0,"",{price_ref})'
    assert book["接单表"]["D4"].value == "SYNTH-OLD-ORDER"
    assert book["正单评审表"]["D5"].value == "SYNTH-OLD-REVIEW-2"


def test_reservation_banner_alone_is_a_boundary_when_no_subtotal_exists():
    book = book_with_sections(subtotal=False)
    result = slots.output_slots(book, 6, top_order_area=True)
    for name in NAMES:
        assert result[name] == list(range(4, 10))
        assert book[name]["A10"].value == "备料单"
        assert row_with_value(book[name], f"SYNTH-RESERVATION-{name}") == 13


@pytest.mark.parametrize("occupied, expected", [
    ((), [4, 5]),
    ((4,), [5, 6]),
    ((4, 6), [7, 8]),
    ((4, 5, 6), [7, 8]),
])
def test_trailing_capacity_or_insertion_never_reuses_an_interior_gap(occupied, expected):
    book = book_with_sections((7, 7, 7))
    for name in NAMES:
        for number in occupied:
            historical_order(book[name], number, f"SYNTH-HISTORY-{number}")
    result = slots.output_slots(book, 2, top_order_area=True)
    for name in NAMES:
        sheet = book[name]
        assert result[name] == expected
        assert max(result[name]) < row_with_value(sheet, "合计：")
        for number in occupied:
            assert find_row(sheet, f"SYNTH-HISTORY-{number}") == number
        if 6 in occupied and 5 not in occupied:
            assert sheet["D5"].value is None


def test_subtotal_that_previously_omitted_reserved_rows_includes_all_new_orders():
    book = book_with_sections((7, 7, 7))
    for name in NAMES:
        sheet = book[name]
        historical_order(sheet, 4, "SYNTH-OLD-ORDER")
        quantity = "K" if name == "ITEM表" else "J"
        sheet[f"{quantity}7"] = f"=SUM({quantity}4:{quantity}4)"
    result = slots.output_slots(book, 6, top_order_area=True)
    for name in NAMES:
        sheet = book[name]
        quantity = "K" if name == "ITEM表" else "J"
        assert result[name] == list(range(5, 11))
        assert sheet[f"{quantity}11"].value == f"=SUM({quantity}4:{quantity}10)"


@pytest.mark.parametrize("word", ["备料单", "合计", "已走货订单", "取消单"])
def test_real_order_type_and_remark_words_do_not_become_section_boundaries(word):
    book = book_with_sections((8, 8, 8))
    for name in NAMES:
        sheet = book[name]
        historical_order(sheet, 4, "SYNTH-ACTUAL-ORDER")
        sheet["B4"] = word
        # A user-owned note without order identity must remain occupied, but a
        # keyword in the far remarks column is not a layout banner.
        sheet.cell(5, 30 if name == "ITEM表" else 20, word)
    result = slots.output_slots(book, 1, top_order_area=True)
    for name in NAMES:
        assert result[name] == [6]
        assert book[name]["B4"].value == word
        assert row_with_value(book[name], "合计：") == 8


def test_inserted_rows_keep_detail_style_and_shift_protected_structures_and_formulas():
    book = book_with_sections()
    sheet = book["ITEM表"]
    sheet.row_dimensions[4].height = 36
    sheet.row_dimensions[6].height = 24
    sheet.merge_cells("A6:H6")
    sheet.merge_cells("A10:H10")
    sheet.print_area = "A1:AB15"
    sheet.print_title_rows = "1:3"
    sheet.auto_filter.ref = "A3:AB4"
    sheet.freeze_panes = "I4"
    line = Side(style="thin", color="123456")
    for cell in sheet[4]:
        cell.fill = PatternFill("solid", fgColor="DBEEDD")
        cell.font = Font(name="Arial", size=11, color="123456")
        cell.border = Border(left=line, right=line, top=line, bottom=line)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    sheet["M4"] = "=IFERROR(K4/L4,\"\")"
    style = copy(sheet["M4"]._style)
    validation = DataValidation(type="list", formula1='"正单,备料单"', allow_blank=True)
    validation.add("B4")
    validation.add("B8")
    sheet.add_data_validation(validation)
    sheet["M8"] = "=K8/L8"
    extra = book.create_sheet("用户资料")
    extra["A1"] = "SYNTH-PRESERVE"
    extra["B1"] = "='ITEM表'!K5"
    extra["C1"] = "='ITEM表'!K8"
    result = slots.output_slots(book, 6, top_order_area=True)
    book = reopened(book)
    sheet = book["ITEM表"]
    assert result["ITEM表"] == list(range(4, 10))
    for number in range(4, 10):
        assert sheet.cell(number, 13).value == f'=IFERROR(K{number}/L{number},"")'
        assert sheet.cell(number, 13)._style == style
        assert sheet.row_dimensions[number].height == 36
        assert any(f"B{number}" in check for check in sheet.data_validations.dataValidation)
    assert sheet["K10"].value == "=SUM(K4:K9)"
    assert sheet["M13"].value == "=K13/L13"
    assert sheet.row_dimensions[11].height == 24
    assert {"A11:H11", "A15:H15"} <= {str(area) for area in sheet.merged_cells.ranges}
    assert sheet.print_area == "'ITEM表'!$A$1:$AB$20"
    assert sheet.print_title_rows == "$1:$3"
    assert sheet.auto_filter.ref == "A3:AB9"
    assert sheet.freeze_panes == "I4"
    assert book["用户资料"]["A1"].value == "SYNTH-PRESERVE"
    assert book["用户资料"]["B1"].value == "='ITEM表'!K10"
    assert book["用户资料"]["C1"].value == "='ITEM表'!K13"


def test_inserted_detail_rows_translate_array_formulas_without_copying_the_subtotal():
    book = book_with_sections()
    book["ITEM表"]["M4"] = ArrayFormula(ref="M4", text="=IFERROR(K4/L4,0)")
    result = slots.output_slots(book, 6, top_order_area=True)
    sheet = reopened(book)["ITEM表"]
    for number in result["ITEM表"]:
        value = sheet.cell(number, 13).value
        assert isinstance(value, ArrayFormula)
        assert value.ref == f"M{number}"
        assert value.text == f"=IFERROR(K{number}/L{number},0)"
    assert sheet["K10"].value == "=SUM(K4:K9)"


def test_section_conditional_formats_follow_rows_and_keep_earlier_ranges_unchanged():
    book = book_with_sections()
    sheet = book["ITEM表"]
    sheet.conditional_formatting.add("F8:H8", FormulaRule(formula=['$E8<>""']))
    sheet.conditional_formatting.add("J8:L8", FormulaRule(
        formula=['AND($E$8<>"",SUM($K$4:$K$4)>0)'],
    ))
    sheet.conditional_formatting.add("A4:B4", FormulaRule(
        formula=["COUNTIF($E$4:$E$4,E4)>0"],
    ))
    extra = book.create_sheet("用户资料")
    extra.conditional_formatting.add("A2:A3", FormulaRule(
        formula=['\'ITEM表\'!$E$8<>""'],
    ))
    slots.output_slots(book, 6, sheet_counts={"ITEM表": 6}, top_order_area=True)
    book = reopened(book)
    actual = {
        str(region.sqref): [rule.formula for rule in rules]
        for region, rules in book["ITEM表"].conditional_formatting._cf_rules.items()
    }
    assert actual == {
        "F13:H13": [['$E13<>""']],
        "J13:L13": [['AND($E$13<>"",SUM($K$4:$K$4)>0)']],
        "A4:B4": [["COUNTIF($E$4:$E$4,E4)>0"]],
    }
    extra_rules = list(book["用户资料"].conditional_formatting._cf_rules.items())
    assert str(extra_rules[0][0].sqref) == "A2:A3"
    assert extra_rules[0][1][0].formula == ['\'ITEM表\'!$E$13<>""']


def test_section_validation_references_move_but_literal_lists_and_earlier_ranges_do_not():
    book = book_with_sections()
    sheet = book["ITEM表"]
    literal = DataValidation(type="list", formula1='"正单,备料单"', allow_blank=True)
    literal.add("B8")
    sheet.add_data_validation(literal)
    numeric = DataValidation(type="whole", operator="between", formula1="=$K$8", formula2="=$L8")
    numeric.add("C8")
    sheet.add_data_validation(numeric)
    earlier = DataValidation(type="custom", formula1="=SUM($K$4:$K$4)>0")
    earlier.add("D8")
    sheet.add_data_validation(earlier)
    extra = book.create_sheet("用户资料")
    linked = DataValidation(type="list", formula1="='ITEM表'!$E$8:$E$9")
    linked.add("A2:A3")
    extra.add_data_validation(linked)
    slots.output_slots(book, 6, sheet_counts={"ITEM表": 6}, top_order_area=True)
    book = reopened(book)
    actual = {
        str(validation.sqref): (validation.formula1, validation.formula2)
        for validation in book["ITEM表"].data_validations.dataValidation
    }
    assert actual == {
        "B13": ('"正单,备料单"', None),
        "C13": ("=$K$13", "=$L13"),
        "D13": ("=SUM($K$4:$K$4)>0", None),
    }
    linked = book["用户资料"].data_validations.dataValidation[0]
    assert str(linked.sqref) == "A2:A3"
    assert linked.formula1 == "='ITEM表'!$E$13:$E$14"


def test_default_slot_mode_does_not_use_the_shushupapa_rule_migration_wrapper(monkeypatch):
    from app.services import customer_order_top_area

    def unexpected(*args, **kwargs):
        raise AssertionError("Other export modes must retain their established inserter")

    monkeypatch.setattr(customer_order_top_area, "insert_rows", unexpected)
    book = book_with_sections()
    result = slots.output_slots(book, 6)
    for name in NAMES:
        assert len(result[name]) == 6
        assert max(result[name]) < row_with_value(book[name], "取消单")


def test_shared_conditional_rule_and_validation_keep_each_sheet_reference_context():
    book = book_with_sections()
    item = book["ITEM表"]
    extra = book.create_sheet("用户资料")
    expression = "AND($E8>0,'ITEM表'!$K$8>0)"
    shared_rule = FormulaRule(formula=[expression])
    shared_validation = DataValidation(type="custom", formula1="=" + expression)
    shared_validation.add("C8")
    for sheet in (item, extra):
        sheet.conditional_formatting.add("F8:H8", shared_rule)
        sheet.add_data_validation(shared_validation)

    slots.output_slots(book, 6, sheet_counts={"ITEM表": 6}, top_order_area=True)
    book = reopened(book)
    for name, number, expected in (
        ("ITEM表", 13, "AND($E13>0,'ITEM表'!$K$13>0)"),
        ("用户资料", 8, "AND($E8>0,'ITEM表'!$K$13>0)"),
    ):
        sheet = book[name]
        region, rules = next(iter(sheet.conditional_formatting._cf_rules.items()))
        assert str(region.sqref) == f"F{number}:H{number}"
        assert rules[0].formula == [expected]
        validation = sheet.data_validations.dataValidation[0]
        assert str(validation.sqref) == f"C{number}"
        assert validation.formula1 == "=" + expected


@pytest.mark.parametrize("customer", ["maxx", "barter", "360", "disney", "dickie"])
def test_other_huaxing_customers_keep_existing_cancellation_based_slot_routing(customer):
    book = book_with_sections()
    expected = slots.output_slots(reopened(book), 1)
    row = order(_huaxing_customer=customer)
    huaxing.export_rows(book, [row])
    for name in NAMES:
        assert find_row(book[name], row["reference_no"]) == expected[name][0]
