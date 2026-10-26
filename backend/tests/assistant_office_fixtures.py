"""Synthetic real OOXML packages, also used by isolated browser acceptance."""
from datetime import datetime
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree


def rewrite(raw, changes):
    output = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(output, 'w', ZIP_DEFLATED) as target:
        entries = {name: source.read(name) for name in source.namelist()}
        entries.update(changes)
        for name, content in entries.items():
            if content is not None:
                target.writestr(name, content)
    return output.getvalue()


def docx_fixture():
    from docx import Document
    doc = Document()
    for n in range(1, 7):
        doc.add_paragraph('' if n == 3 else f'合成资料第 {n} 段')
    doc.add_heading('验收条件', level=1)
    doc.add_paragraph('合成条款：交付前核对数量与单位。')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '合成表格内容'
    output = BytesIO(); doc.save(output)
    return output.getvalue()


def xlsx_fixture():
    from openpyxl import Workbook
    book = Workbook(); sheet = book.active; sheet.title = '采购'
    sheet.merge_cells('A1:D1'); sheet['A1'] = '合成采购数量（件）'
    sheet['A2'] = '00017'; sheet['A3'] = datetime(2026, 10, 23)
    sheet['A3'].number_format = 'yyyy-mm-dd'
    sheet['A7'] = 500; sheet['B7'] = 100; sheet['C7'] = '=A7-B7'
    sheet['D7'] = '=EXTERNAL_REF'; sheet['E7'] = '=A7-A7'
    sheet.row_dimensions[7].hidden = True
    book.create_sheet('隐藏资料')['A1'] = '隐藏表中的合成值'
    book['隐藏资料'].sheet_state = 'hidden'
    output = BytesIO(); book.save(output)
    raw = output.getvalue()
    with ZipFile(BytesIO(raw)) as source:
        root = etree.fromstring(source.read('xl/worksheets/sheet1.xml'))
    ns = {'s': root.nsmap[None]}
    for address, value in [('C7', '400'), ('E7', '0')]:
        root.find(f'.//s:c[@r="{address}"]/s:v', ns).text = value
    return rewrite(raw, {'xl/worksheets/sheet1.xml': etree.tostring(root)})
