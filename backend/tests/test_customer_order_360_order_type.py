"""Huakang A 360 business order types are independent of RL routing types."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO

import openpyxl
import pytest

from app.services import customer_order_unified as service
from app.schemas.customer_order import CustomerOrderLineOut
from app.services.huakang_a_order_legacy import schedule_parser, unified_schedule as mapping
from test_customer_order_huakang_a_unified import template
from customer_order_360_private_samples import private_po_samples


def release_record(source_type):
    return {
        "file_name": "release.pdf", "order_type": source_type,
        "contract_no": "RL-765432-6", "customer_release": "RL76543226",
        "customer_po": "7654322", "item_no": "NEW-360", "quantity": 120,
        "master_carton_qty": 6, "customer_name": "Example Retailer", "description": "Toy",
        "planned_inspection_date": "2027-01-08", "factory_commit_date": "2027-01-08",
        "revision": 2, "parse_ok": True, "parse_warnings": [],
    }


@pytest.mark.parametrize("source_type", ["DI Order", "Domestic Order", "", None])
def test_new_360_release_is_normal_order_with_source_type_preserved(source_type):
    source = release_record(source_type)
    original = deepcopy(source)
    row = mapping.map_record({"order_type": source_type}, source, "360")
    assert row["order_type"] == "正单"
    assert "source_order_type" not in row
    assert row["lineage"]["source_order_type"] == f"release.pdf · RL Order Type: {source_type or '未提供'}"
    assert "新 PURCHASE ORDER RELEASE 记为正单" in row["lineage"]["order_type"]
    assert source == original


@pytest.mark.parametrize("customer", ["green-toys", "headstart"])
def test_other_huakang_a_customer_order_type_is_unchanged(customer):
    row = mapping.map_record({"order_type": "加单"}, release_record("DI Order"), customer)
    assert row["order_type"] == "加单"
    assert "source_order_type" not in row
    assert "order_type" not in row.get("lineage", {})
    assert not mapping.enabled("huaxing", "360")


def test_new_360_type_is_written_to_item_and_linked_summaries_without_retyping_history(monkeypatch):
    source = release_record("DI Order")
    monkeypatch.setattr(schedule_parser, "parse_po_file", lambda *_: deepcopy(source))
    original = openpyxl.load_workbook(BytesIO(template("360")))
    old_types = {service.ITEM_SHEET: "备料单", service.ORDER_SHEET: "补单", service.REVIEW_SHEET: "改单"}
    old_products = {name: original[name].cell(4, 8 if name == service.ITEM_SHEET else 7).value
                    for name in old_types}
    for name, order_type in old_types.items():
        original[name]["B4"] = order_type
    original_bytes = BytesIO()
    original.save(original_bytes)
    output, _, preview = service.export_unified_customer_schedule(
        customer_code="360", factory_id="huakang-a", received_date="2026-10-19",
        po_files=[("release.pdf", b"synthetic release")], schedule_file_name="schedule.xlsx",
        schedule_content=original_bytes.getvalue(),
    )
    assert preview["rows"][0]["order_type"] == "正单"
    assert "source_order_type" not in preview["rows"][0]
    assert preview["rows"][0]["lineage"]["source_order_type"] == "release.pdf · RL Order Type: DI Order"
    serialized = CustomerOrderLineOut.model_validate(preview["rows"][0]).model_dump()
    assert serialized["order_type"] == "正单"
    workbook = openpyxl.load_workbook(BytesIO(output))
    item = workbook[service.ITEM_SHEET]
    item_row = next(c.row for c in item._cells.values() if c.column == 8 and c.value == "NEW-360")
    assert item.cell(item_row, 2).value == "正单"
    for name in (service.ORDER_SHEET, service.REVIEW_SHEET):
        sheet = workbook[name]
        summary_row = next(c.row for c in sheet._cells.values() if c.column == 7 and c.value == "NEW-360")
        assert sheet.cell(summary_row, 2).value == f'=IF(LEN(\'ITEM表\'!B{item_row})=0,"",\'ITEM表\'!B{item_row})'
    for name, order_type in old_types.items():
        product_column = 8 if name == service.ITEM_SHEET else 7
        product = old_products[name]
        old_row = next(c.row for c in workbook[name]._cells.values() if c.column == product_column and c.value == product)
        assert workbook[name].cell(old_row, 2).value == order_type


@pytest.mark.parametrize("sample", private_po_samples())
def test_supplied_po_source_type_is_preserved_while_new_order_is_normal(sample):
    source, source_type = sample["source"], sample["expected"]["order_type"]
    before = sha256(source.read_bytes()).hexdigest()
    record = schedule_parser.parse_po_file(source, sample["manifest"]["po_directory"])
    original = deepcopy(record)
    row = mapping.map_record({}, record, "360")
    assert record["order_type"] == source_type
    assert row["order_type"] == "正单" and "source_order_type" not in row
    assert row["lineage"]["source_order_type"] == f"{source.name} · RL Order Type: {source_type or '未提供'}"
    assert record == original
    assert sha256(source.read_bytes()).hexdigest() == before
