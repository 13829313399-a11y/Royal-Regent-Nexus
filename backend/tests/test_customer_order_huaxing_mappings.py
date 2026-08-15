from copy import copy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

from app.services.customer_order_huaxing import (
    HUAXING_CUSTOMER_MAPPINGS,
    HuaxingCustomerOrderError,
    PreparedBatch,
    _dedupe_multi_orders,
    _export_prepared,
    _issues,
    _record_fields,
    _validate_skips,
    create_huaxing_customer_preview,
    export_huaxing_customer_schedule,
)
from app.services.customer_order_manual import (
    apply_overrides_to_records,
    decorate_manual_resolution_policy,
)
from app.services.huaxing_order_legacy import (
    edu_schedule,
    multi_schedule,
    new_order_excel,
    shixin_schedule,
    three_sixty_schedule,
    yinhui_po_parser,
    yinhui_schedule,
)


def test_huaxing_customer_order_center_exposes_exactly_six_new_mappings():
    assert set(HUAXING_CUSTOMER_MAPPINGS) == {
        "edu",
        "360",
        "yinhui",
        "seasons",
        "maxx",
        "shushupapa",
    }
    assert all(spec.target_template.endswith("_SCHEDULE_APPEND_V2") for spec in HUAXING_CUSTOMER_MAPPINGS.values())


def test_edu_inspection_date_is_seven_days_before_ship_date_and_avoids_weekend():
    row = {
        "customer_type": "EDU",
        "customer_po": "EDU-100",
        "item_no": "A100",
        "product_name": "Test",
        "quantity": 120,
        "case_pack": 12,
        "ship_date": "2026-08-16",
    }

    edu_schedule.add_derived_fields(row, today=date(2026, 8, 1))

    assert row["inspection_date"] == "2026-08-07"
    assert row["cartons"] == 10


def test_edu_export_synchronizes_item_review_and_order_sheets(tmp_path: Path):
    workbook = Workbook()
    item = workbook.active
    item.title = "Iteam表"
    item_headers = [
        "客出单日期", "客人PO", "华兴PO", "是否已入系统", "主合同号",
        "SO.NO", "客名", "国家", "产品编号", "客货号", "产品名称",
        "数量", "装箱", "包装", "生产日期码", "客要求走货期",
    ]
    for column, value in enumerate(item_headers, 1):
        item.cell(3, column, value)
        item.cell(4, column).font = Font(color="000000")
    item["B4"] = "PO-OLD"
    item["C4"] = "EDUHX00001"
    item["I4"] = "ITEM-OLD"
    item["L4"] = 1

    review = workbook.create_sheet("正单评审表")
    review_headers = [
        "客出单日期", "客人PO", "华兴PO", "主合同号", "SO.NO", "客名",
        "国家", "产品编号", "客货号", "产品名称", "数量", "装箱",
        "生产日期码", "验货期", "走货期", "单价", "金额HKD",
    ]
    for column, value in enumerate(review_headers, 1):
        review.cell(4, column, value)
    review["B5"] = "PO-OLD"
    review["C5"] = "EDUHX00001"
    review["D5"] = "=VLOOKUP(C5,Iteam表!C:E,3,0)"
    review["K5"] = "=VLOOKUP(C5,Iteam表!C:L,10,0)"

    order = workbook.create_sheet("接单表")
    order_headers = [
        "客出单日期", "客人PO", "华兴PO", "主合同号", "SO.NO", "客名",
        "国家", "产品编号", "客货号", "产品名称", "数量", "单价", "金额",
        "客要求走货期",
    ]
    for column, value in enumerate(order_headers, 1):
        order.cell(3, column, value)
    order["B4"] = "PO-OLD"
    order["C4"] = "EDUHX00001"
    order["D4"] = "=VLOOKUP(C4,Iteam表!C:E,3,0)"
    order["K4"] = "=VLOOKUP(C4,Iteam表!C:L,10,0)"
    order["L5"] = "合计"
    order["M5"] = "=SUM(M4:M4)"

    source = BytesIO()
    workbook.save(source)
    workbook.close()
    output_path = tmp_path / "EDU排期_新单.xlsx"
    record = {
        "order_date": "2026-08-14",
        "customer_po": "PO-NEW",
        "huaxing_po": "EDUHX00002",
        "contract_no": "SC-NEW",
        "customer": "EDU UK",
        "country": "英国",
        "item_no": "ITEM-NEW",
        "product_name": "新产品",
        "quantity": 120,
        "case_pack": 12,
        "ship_date": "2026-09-30",
        "unit_price": 10,
        "amount": 1200,
    }

    _export_prepared(
        customer_code="edu",
        prepared=PreparedBatch([record], [], "Iteam表"),
        schedule_file_name="EDU排期.xlsx",
        schedule_content=source.getvalue(),
        output_path=output_path,
    )

    rendered = load_workbook(output_path, data_only=False)
    assert rendered["Iteam表"]["C5"].value == "EDUHX00002"
    assert rendered["Iteam表"]["I5"].value == "ITEM-NEW"
    assert rendered["正单评审表"]["C6"].value == "EDUHX00002"
    assert rendered["正单评审表"]["D6"].value == "=VLOOKUP(C6,Iteam表!C:E,3,0)"
    assert rendered["接单表"]["C5"].value == "EDUHX00002"
    assert rendered["接单表"]["D5"].value == "=VLOOKUP(C5,Iteam表!C:E,3,0)"
    assert rendered["接单表"]["L6"].value == "合计"
    assert rendered["接单表"]["M6"].value == "=SUM(M4:M5)"
    rendered.close()


def test_360_date_code_uses_factory_year_and_day_of_year():
    assert three_sixty_schedule.build_360_date_code("2026-07-02", "60350") == "60350A26183"


def test_360_manual_inspection_result_clears_overdue_blocker_and_is_exportable(tmp_path):
    record = {
        "production_no": "RL-450484",
        "contract_no": "450484",
        "item_no": "1212016551.04",
        "quantity": 100,
        "inspection_date": "2026-08-01",
        "source_sheet": "接单表",
    }
    flags = three_sixty_schedule.build_flags(record, today=date(2026, 8, 7))
    issue = next(flag for flag in flags if flag["code"] == "missing_inspection_result")
    assert issue["field"] == "inspection_result"

    apply_overrides_to_records(
        [record],
        ["360-1"],
        [{
            "row_id": "360-1",
            "issue_key": "360-1:missing_inspection_result",
            "field": "inspection_result",
            "value": "合格",
        }],
    )
    assert record["inspection_result"] == "合格"
    assert not any(
        flag.get("code") == "missing_inspection_result"
        for flag in three_sixty_schedule.build_flags(record, today=date(2026, 8, 7))
    )

    output_path = tmp_path / "360人工补录.xlsx"
    three_sixty_schedule.create_import_workbook([record], output_path)
    workbook = load_workbook(output_path, data_only=True)
    worksheet = workbook["AI导入待确认"]
    headers = {cell.value: cell.column for cell in worksheet[1]}
    assert worksheet.cell(2, headers["验货结果"]).value == "合格"
    workbook.close()


def test_yinhui_uses_fixed_exchange_rate_and_ship_minus_five_days():
    row = {
        "so_no": "SO-1",
        "contract_no": "YH-1",
        "customer": "YINHUI",
        "item_no": "ITEM-1",
        "product_name": "Test",
        "quantity": 10,
        "case_pack": 5,
        "po_ship_date": "2026-08-20",
        "unit_price_usd": 2,
    }

    yinhui_schedule.add_derived_fields(row, today=date(2026, 8, 1))

    assert row["unit_price_hkd"] == 15.5
    assert row["total_hkd"] == 155
    assert row["inspection_date"] == "2026-08-15"


def _yinhui_ocr_text(*, ship_date: str = "2026.12.06", declared_pages: int = 2) -> str:
    return f"""
Page: 1 of {declared_pages}
PURCHASE ORDER: 4500002406
Unloading Point: SWEDEN Date (YMD): 2026.06.16
ITEM# PART# DESCRIPTION DELI.(YMD) QTY UNIT PRICE AMOUNT
00010 5122208863801 RESCUE ELLI BANANA {ship_date} 20 PCS 0.1000 2.00
Packaging: POLYBAG PR:3000010075
S0O:2300000463 -780
Packing (Inner/Outer): 0/0
Name:SILVERLIT NORDIC AB
00020 5122308863801 RESCUE ELLI TEETH {ship_date} 20 PCS 0.1000 2.00
Packaging: POLYBAG PR:3000010081
SO:2300000463 -790
Packing (Inner/Outer): 0/0
Name:SILVERLIT NORDIC AB
Total: 4.00
**SAY: USD FOUR DOLLARS AND ZERO CENTS ONLY**
""".strip()


def _make_yinhui_three_sheet_template() -> Workbook:
    workbook = Workbook()
    item = workbook.active
    item.title = "Iteam表"
    item_headers = [
        "出单日期", "SO", "银辉合同号", "客名", "產品編號", "产品名称",
        "PO数量", "装箱数量", "外箱装箱数", "说明书", "彩盒", "客贴纸",
        "贴位图", "日期码", "备注", "箱唛资料", "验货日期", "完成情况",
        "备注", "订单单价USD", "单价HK$", "总金额HK$", "总金额USD",
        "总金额HK$", "出厂价HK$", "出厂价总金额HK$", "备注", "系统",
        "MS Container Type (FCL/LCL)", "Consolidated ID", "Port of Discharge",
        "走货日期", "车次", "走货方式", "发票单价", "发票金额", "打票日期",
        "发票号码", "其他", "其他金额",
    ]
    for column, value in enumerate(item_headers, 1):
        item.cell(4, column, value)
        item.column_dimensions[item.cell(4, column).column_letter].width = 16
    item_history = [
        "2026-07-01", "SO-HISTORY", "PO-HISTORY", "历史客户", "ITEM-HISTORY", "历史产品",
        10, 5, "=G5/H5", "说明书", "彩盒", "贴纸", None, "180B6RR", "历史备注", "箱唛",
        "2026-07-20", None, None, 2, "=T5*7.75", "=U5*G5", "=T5*G5",
        "=U5*G5", "=T5*7.75", "=Y5*G5", None, None, None, None, None,
        "2026-07-25", None, None, None, None, None, None, None, None,
    ]
    blue = "0000FF"
    for column, value in enumerate(item_history, 1):
        cell = item.cell(5, column, value)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.font = Font(color=blue)
    item.row_dimensions[5].height = 28
    item.merge_cells("E6:F6")
    item["E6"] = "合计"
    item["G6"] = "=SUBTOTAL(9,G5:G5)"
    item.row_dimensions[6].height = 36
    item.auto_filter.ref = "A4:AN5"
    item.print_area = "A1:AN6"
    item.freeze_panes = "A5"

    review = workbook.create_sheet("正单评审表")
    review_headers = [
        "款号", "客名", "来单日期", "主合同号", "產品編號", "产品名称", "PO数量",
        "装箱数量", "装箱", "说明书", "彩盒", "PO利宝", "日期码", "备注", "验货期",
        "单价", "金额HK$", "公證行驗貨日期(在華登廠房)", "生产要求",
    ]
    for column, value in enumerate(review_headers, 1):
        review.cell(4, column, value)
    review_links = {
        2: "D", 3: "A", 4: "C", 5: "E", 6: "F", 7: "G", 8: "H", 9: "I",
        10: "J", 11: "K", 12: "L", 13: "N", 14: "O", 15: "Q", 16: "U", 17: "V",
    }
    for column, source_column in review_links.items():
        review.cell(5, column, f"='Iteam表'!{source_column}5")
        review.cell(5, column).font = Font(color=blue)
    review["E10"] = "='Iteam表'!E6"
    review["G10"] = "='Iteam表'!G6"

    order = workbook.create_sheet("接单表")
    order_headers = [
        "来单日期", "主合同号", "客名", "產品型號", "产品名称", "PO数量", "裝箱隻數",
        "总箱数", "单价HKD", "金额HKD", "彩盒", "PO利宝", "备注", "验货期", "日期码",
        "验货结果", "第三方公证行验货",
    ]
    for column, value in enumerate(order_headers, 1):
        order.cell(4, column, value)
    order_links = {
        1: "A", 2: "C", 3: "D", 4: "E", 5: "F", 6: "G", 7: "H", 8: "I",
        9: "U", 10: "V", 11: "K", 12: "L", 13: "O", 14: "Q", 15: "N",
    }
    for column, source_column in order_links.items():
        order.cell(5, column, f"='Iteam表'!{source_column}5")
        order.cell(5, column).font = Font(color=blue)
    order["D10"] = "='Iteam表'!E6"
    order["F10"] = "='Iteam表'!G6"
    return workbook


def test_yinhui_scanned_pdf_ocr_extracts_rows_and_keeps_zero_pack_missing(monkeypatch):
    monkeypatch.setattr(
        yinhui_po_parser,
        "_pdf_text",
        lambda _source: (_yinhui_ocr_text(declared_pages=3), 2, True),
    )

    parsed = yinhui_po_parser.parse_pdf(b"%PDF-scan", "RR-4500002406.pdf")

    assert parsed["meta"] == {
        "pages": 2,
        "text_chars": len(_yinhui_ocr_text(declared_pages=3)),
        "po_no": "4500002406",
        "used_ocr": True,
    }
    assert len(parsed["rows"]) == 2
    first, second = parsed["rows"]
    assert first["so_no"] == second["so_no"] == "2300000463"
    assert first["customer"] == "SILVERLIT NORDIC AB"
    assert first["item_no"] == "5122208863801"
    assert second["item_no"] == "5122308863801"
    assert first["quantity"] == second["quantity"] == 20
    assert first["unit_price_usd"] == second["unit_price_usd"] == 0.1
    assert first["total_usd"] == second["total_usd"] == 2
    assert first["case_pack"] is None
    assert first["cartons"] is None
    assert first["inspection_date"] == "2026-12-01"
    assert any(flag["code"] == "missing_case_pack" for flag in first["flags"])
    assert any("实际为 2 页" in warning for warning in parsed["warnings"])


def test_yinhui_scanned_pdf_manual_pack_override_is_written_to_export(monkeypatch):
    monkeypatch.setattr(
        yinhui_po_parser,
        "_pdf_text",
        lambda _source: (_yinhui_ocr_text(), 2, True),
    )
    workbook = _make_yinhui_three_sheet_template()
    schedule_stream = BytesIO()
    workbook.save(schedule_stream)
    workbook.close()
    schedule_content = schedule_stream.getvalue()

    preview = create_huaxing_customer_preview(
        customer_code="yinhui",
        factory_id="huaxing",
        received_date="2026-08-11",
        po_files=[("RR-4500002406.pdf", b"%PDF-scan")],
        schedule_file_name="2026-银辉排期.xlsx",
        schedule_content=schedule_content,
    )
    decorate_manual_resolution_policy(preview)
    pack_issues = [
        next(issue for issue in row["issues"] if issue["code"] == "missing_case_pack")
        for row in preview["rows"]
    ]
    assert preview["summary"] == {"total": 2, "valid": 0, "warning": 0, "blocked": 2}
    assert all(issue["can_edit"] for issue in pack_issues)
    assert all(issue["edit_field"] == "case_pack" for issue in pack_issues)
    assert all(issue["edit_input_type"] == "number" for issue in pack_issues)

    overrides = [
        {
            "row_id": row["id"],
            "issue_key": issue["skip_key"],
            "field": "case_pack",
            "value": "4",
        }
        for row, issue in zip(preview["rows"], pack_issues, strict=True)
    ]
    output, _, _ = export_huaxing_customer_schedule(
        customer_code="yinhui",
        factory_id="huaxing",
        received_date="2026-08-11",
        po_files=[("RR-4500002406.pdf", b"%PDF-scan")],
        schedule_file_name="2026-银辉排期.xlsx",
        schedule_content=schedule_content,
        skipped_issue_keys={issue["skip_key"] for issue in pack_issues},
        manual_overrides=overrides,
    )

    rendered = load_workbook(BytesIO(output), data_only=False)
    target = rendered["Iteam表"]
    assert target["B6"].value == target["B7"].value == "2300000463"
    assert target["C6"].value == target["C7"].value == "4500002406"
    assert target["E6"].value == "5122208863801"
    assert target["E7"].value == "5122308863801"
    assert target["H6"].value == target["H7"].value == 4
    assert target["I6"].value == "=IFERROR(G6/H6,5.0)"
    assert target["I7"].value == "=IFERROR(G7/H7,5.0)"
    assert target["V6"].value == "=G6*U6"
    assert target["W6"].value == "=G6*T6"
    assert target["X6"].value == "=G6*U6"
    assert target["Y6"].value == "=T6*7.75"
    assert target["Z6"].value == "=Y6*G6"
    assert target["AF6"].value == "2026-12-06"
    assert target["E8"].value == "合计"
    assert target["G8"].value == "=SUBTOTAL(9,G5:G7)"

    review = rendered["正单评审表"]
    assert review["D6"].value == "='Iteam表'!C6"
    assert review["D7"].value == "='Iteam表'!C7"
    assert review["Q6"].value == "='Iteam表'!V6"
    assert review["E10"].value == "='Iteam表'!E8"

    order = rendered["接单表"]
    assert order["B6"].value == "='Iteam表'!C6"
    assert order["B7"].value == "='Iteam表'!C7"
    assert order["J6"].value == "='Iteam表'!V6"
    assert order["D10"].value == "='Iteam表'!E8"
    rendered.close()


def test_yinhui_export_preserves_complete_workbook_and_inserts_before_total(tmp_path: Path):
    source_path = tmp_path / "银辉原排期.xlsx"
    output_path = tmp_path / "银辉原排期_银辉新单.xlsx"
    image_path = tmp_path / "marker.png"
    Image.new("RGB", (8, 8), "red").save(image_path)

    workbook = _make_yinhui_three_sheet_template()
    linked = workbook["接单表"]
    linked["A1"] = "='Iteam表'!G6"
    target = workbook["Iteam表"]
    image_sheet = workbook.create_sheet("历史图片")
    image_sheet["A1"] = "历史资料"
    image_sheet.add_image(SpreadsheetImage(image_path), "B2")
    target.page_setup.orientation = "landscape"
    target.conditional_formatting.add(
        "E6:F6",
        CellIsRule(operator="equal", formula=['"合计"'], fill=PatternFill("solid", fgColor="FFF2CC")),
    )
    validation = DataValidation(type="list", formula1='"A,B"')
    validation.add("E6")
    target.add_data_validation(validation)
    workbook.save(source_path)
    workbook.close()
    original_bytes = source_path.read_bytes()

    yinhui_schedule.create_export(
        [{
            "order_date": "2026-08-07",
            "so_no": "SO-NEW",
            "contract_no": "PO-NEW",
            "customer": "新客户",
            "item_no": "ITEM-NEW",
            "product_name": "新产品",
            "quantity": 20,
            "case_pack": 4,
            "unit_price_usd": 3,
            "factory_unit_price_hkd": 20,
        }],
        output_path,
        source_path,
        template_filename=source_path.name,
    )

    assert source_path.read_bytes() == original_bytes
    rendered = load_workbook(output_path, data_only=False)
    assert rendered.sheetnames == ["Iteam表", "正单评审表", "接单表", "历史图片"]
    assert rendered["历史图片"]["A1"].value == "历史资料"
    assert len(rendered["历史图片"]._images) == 1
    assert rendered["接单表"]["A1"].value == "='Iteam表'!G7"
    target = rendered["Iteam表"]
    assert target["B5"].value == "SO-HISTORY"
    assert target["B6"].value == "SO-NEW"
    assert target["I6"].value == "=IFERROR(G6/H6,5.0)"
    assert target["U6"].value == 23.25
    assert target["X6"].value == "=G6*U6"
    assert target["Y6"].value == 20
    assert target["Z6"].value == 400
    assert target["E7"].value == "合计"
    assert target["G7"].value == "=SUBTOTAL(9,G5:G6)"
    assert "E7:F7" in {str(value) for value in target.merged_cells.ranges}
    assert target.row_dimensions[6].height == 28
    assert target.row_dimensions[7].height == 36
    assert target["A6"].fill.fgColor.rgb == target["A5"].fill.fgColor.rgb
    assert target["A6"]._style.fontId == target["A5"]._style.fontId
    assert target.auto_filter.ref == "A4:AN6"
    assert str(target.print_area).endswith("$A$1:$AN$7")
    assert target.page_setup.orientation == "landscape"
    assert target.freeze_panes == "A5"
    assert "E7:F7" in {str(value.sqref) for value in target.conditional_formatting}
    assert {str(value.sqref) for value in target.data_validations.dataValidation} == {"E7"}
    assert rendered["正单评审表"]["D6"].value == "='Iteam表'!C6"
    assert rendered["正单评审表"]["E10"].value == "='Iteam表'!E7"
    assert rendered["接单表"]["B6"].value == "='Iteam表'!C6"
    assert rendered["接单表"]["D10"].value == "='Iteam表'!E7"
    rendered.close()


def test_yinhui_export_keeps_source_row_font_color_instead_of_forcing_blue(tmp_path: Path):
    source_path = tmp_path / "银辉黑字排期.xlsx"
    output_path = tmp_path / "银辉黑字排期_新单.xlsx"
    workbook = _make_yinhui_three_sheet_template()
    for worksheet in workbook.worksheets:
        for cell in worksheet[5]:
            font = Font(
                name=cell.font.name,
                size=cell.font.size,
                bold=cell.font.bold,
                italic=cell.font.italic,
                color="000000",
            )
            cell.font = font
    workbook.save(source_path)
    workbook.close()

    yinhui_schedule.create_export(
        [{
            "order_date": "2026-08-14",
            "so_no": "SO-NEW",
            "contract_no": "PO-NEW",
            "customer": "新客户",
            "item_no": "ITEM-NEW",
            "product_name": "新产品",
            "quantity": 20,
            "case_pack": 4,
            "unit_price_usd": 3,
        }],
        output_path,
        source_path,
        template_filename=source_path.name,
    )

    rendered = load_workbook(output_path, data_only=False)
    source_color = rendered["Iteam表"]["C5"].font.color
    new_color = rendered["Iteam表"]["C6"].font.color
    assert source_color.type == new_color.type == "rgb"
    assert source_color.rgb == new_color.rgb
    rendered.close()


def test_yinhui_export_uses_complete_ordinary_style_when_last_order_is_blue(tmp_path: Path):
    source_path = tmp_path / "银辉末行蓝字排期.xlsx"
    output_path = tmp_path / "银辉末行蓝字排期_新单.xlsx"
    workbook = _make_yinhui_three_sheet_template()
    for worksheet in workbook.worksheets:
        worksheet.row_dimensions[5].height = 31.5
        for column in range(1, worksheet.max_column + 1):
            ordinary = worksheet.cell(5, column)
            ordinary.number_format = (
                "0.000" if column % 2 else "@"
            )
            ordinary.font = copy(ordinary.font)
            ordinary.font = Font(
                name=ordinary.font.name,
                size=ordinary.font.size,
                bold=ordinary.font.bold,
                italic=ordinary.font.italic,
                color="000000",
            )
        if worksheet.title == "Iteam表":
            worksheet.unmerge_cells("E6:F6")
        worksheet.insert_rows(6)
        if worksheet.title == "Iteam表":
            worksheet.merge_cells("E7:F7")
        worksheet.row_dimensions[6].height = worksheet.row_dimensions[5].height
        for column in range(1, worksheet.max_column + 1):
            source = worksheet.cell(5, column)
            destination = worksheet.cell(6, column)
            destination.value = source.value
            destination._style = copy(source._style)
            destination.font = copy(source.font)
            destination.font = Font(
                name=destination.font.name,
                size=destination.font.size,
                bold=destination.font.bold,
                italic=destination.font.italic,
                color="0000FF",
            )
    workbook.save(source_path)
    workbook.close()

    yinhui_schedule.create_export(
        [{
            "order_date": "2026-08-14",
            "so_no": "SO-NEW",
            "contract_no": "PO-NEW",
            "customer": "新客户",
            "item_no": "ITEM-NEW",
            "product_name": "新产品",
            "quantity": 20,
            "case_pack": 4,
            "unit_price_usd": 3,
        }],
        output_path,
        source_path,
        template_filename=source_path.name,
    )

    rendered = load_workbook(output_path, data_only=False)
    for worksheet in rendered.worksheets:
        assert worksheet["C6"].font.color.type == "rgb"
        assert worksheet["C6"].font.color.rgb.endswith("0000FF")
        assert worksheet.row_dimensions[7].height == worksheet.row_dimensions[5].height
        for column in range(1, worksheet.max_column + 1):
            ordinary = worksheet.cell(5, column)
            inserted = worksheet.cell(7, column)
            assert inserted.number_format == ordinary.number_format, (
                f"{worksheet.title}!{inserted.coordinate}: "
                f"{inserted.number_format!r} != {ordinary.number_format!r}"
            )
            assert copy(inserted.font) == copy(ordinary.font)
            assert copy(inserted.fill) == copy(ordinary.fill)
            assert copy(inserted.border) == copy(ordinary.border)
            assert copy(inserted.alignment) == copy(ordinary.alignment)
            assert copy(inserted.protection) == copy(ordinary.protection)
    rendered.close()


def test_360_export_writes_item_review_and_order_sheets_above_totals(tmp_path: Path):
    template_path = tmp_path / "360排期.xlsx"
    output_path = tmp_path / "360排期_360新单.xlsx"
    workbook = Workbook()
    order = workbook.active
    order.title = "接单表"
    review = workbook.create_sheet("正单评审表 ")
    item = workbook.create_sheet("Iteam表")
    workbook.create_sheet("历史资料")["A1"] = "必须保留"

    item_headers = [
        None, "出单日期", "美国PO号", "美国生产单号", "360生产单号", "WM PO号",
        "360合同号", "客名", "產品編號", "产品名称", "规格", "PO数量",
        "外箱装箱数", "箱数", "说明书", "彩盒", "车款", "外箱PO利宝", "日期码",
        "备注", "验货日期", "FCD期", "走货期", "验货结果", None, None, None,
        None, "走货国家", "跟单", "条码", "订单单价USD", "单价HK$", "模具平摊费用",
        "测试平摊费用", "订单单价USD", "单价HK$", "总金额HK$", "总金额USD",
        "总金额HK$", "出厂价HK$", "出厂价总金额HK$",
    ]
    for column, value in enumerate(item_headers, 1):
        item.cell(3, column, value)
    old_item = [
        None, datetime(2026, 7, 1), "PO-OLD", "REL-OLD", "RL-OLD", None,
        "CONTRACT-OLD", "历史客户", "10001", "历史产品", "美国", 10, 2,
        "=L4/M4", "EN", "彩盒", "车款", "利宝", "60350A26183", None,
        datetime(2026, 7, 20), None, datetime(2026, 7, 30), None,
    ]
    for column, value in enumerate(old_item, 1):
        item.cell(4, column, value)
    for coordinate, formula in {
        "AG4": "=AF4*7.75", "AJ4": "=AF4-AH4-AI4", "AK4": "=AJ4*7.75",
        "AL4": "=L4*AK4", "AM4": "=AF4*L4", "AN4": "=AG4*L4", "AP4": "=AO4*L4",
    }.items():
        item[coordinate] = formula
    item["AF4"] = 12
    item["AH4"] = "=100/1000"
    item["AI4"] = "=50/1000"
    item["AO4"] = 90
    item["A4"].fill = PatternFill("solid", fgColor="DDEBF7")
    item["I6"] = "合计"
    item["L6"] = "=SUBTOTAL(9,L4:L5)"

    review_headers = [
        "款号", "客名", "来单日期", "主合同号", "360生产单号", "產品編號",
        "产品名称", "规格", "PO数量", "装箱", "说明书", "彩盒", "车款",
        "PO利宝", "日期码", "备注", "验货期", "单价", "金额HK$",
        "公證行驗貨日期(在華登廠房)", "生产要求",
    ]
    order_headers = [
        "来单日期", "主合同号", "360生产单号", "客名", "產品型號", "产品名称",
        "產品规格", "PO数量", "裝箱隻數", "总箱数", "单价HKD", "金额HKD",
        "彩盒", "PO利宝", "备注", "验货期", "日期码", "验货结果", "第三方公证行验货",
    ]
    for column, value in enumerate(review_headers, 1):
        review.cell(4, column, value)
    for column, value in enumerate(order_headers, 1):
        order.cell(4, column, value)
    for column, formula in {
        3: "=Iteam表!B4", 4: "=Iteam表!G4", 5: "=Iteam表!E4", 6: "=Iteam表!I4",
        7: "=Iteam表!J4", 8: "=Iteam表!K4", 9: "=Iteam表!L4", 10: "=Iteam表!M4",
        11: "=Iteam表!O4", 12: "=Iteam表!P4", 13: "=Iteam表!Q4", 14: "=Iteam表!R4",
        15: "=Iteam表!S4", 16: "=Iteam表!T4", 17: "=Iteam表!U4", 18: "=Iteam表!AG4",
        19: "=I5*R5",
    }.items():
        review.cell(5, column, formula)
    review["F7"] = "=Iteam表!I6"
    for column, formula in {
        1: "=Iteam表!B4", 2: "=Iteam表!G4", 3: "=Iteam表!E4", 4: "=Iteam表!H4",
        5: "=Iteam表!I4", 6: "=Iteam表!J4", 7: "=Iteam表!K4", 8: "=Iteam表!L4",
        9: "=Iteam表!M4", 10: "=H5/I5", 11: "=Iteam表!AG4", 12: "=K5*H5",
        13: "=Iteam表!P4", 14: "=Iteam表!R4", 15: "=Iteam表!T4", 16: "=Iteam表!U4",
        17: "=Iteam表!S4",
    }.items():
        order.cell(5, column, formula)
    order["E7"] = "=Iteam表!I6"
    workbook.save(template_path)
    workbook.close()

    three_sixty_schedule.create_schedule_review_workbook(
        template_path.read_bytes(),
        [{
            "order_date": "2026-08-07",
            "po_no": "PO-NEW",
            "customer_release_no": "REL-NEW",
            "production_no": "RL-NEW",
            "contract_no": "CONTRACT-NEW",
            "customer": "新客户",
            "item_no": "20002",
            "product_name": "新产品",
            "spec": "美国",
            "quantity": 20,
            "outer_pack": 4,
            "inspection_date": "2026-08-20",
            "unit_price_usd": 15,
        }, {
            "order_date": "2026-08-08",
            "po_no": "PO-NEW-2",
            "customer_release_no": "REL-NEW-2",
            "production_no": "RL-NEW-2",
            "contract_no": "CONTRACT-NEW-2",
            "customer": "新客户",
            "item_no": "20003",
            "product_name": "新产品二",
            "spec": "美国",
            "quantity": 8,
            "outer_pack": 4,
            "inspection_date": "2026-08-21",
            "unit_price_usd": 15,
        }],
        output_path,
        template_filename=template_path.name,
    )

    rendered = load_workbook(output_path, data_only=False)
    assert rendered.sheetnames == ["接单表", "正单评审表 ", "Iteam表", "历史资料"]
    item = rendered["Iteam表"]
    assert item["C5"].value == "PO-NEW"
    assert item["E5"].value == "RL-NEW"
    assert item["I5"].value == "20002"
    assert item["C6"].value == "PO-NEW-2"
    assert item["I6"].value == "20003"
    assert item["N5"].value == "=L5/M5"
    assert item["AG5"].value == "=AF5*7.75"
    assert item["AJ5"].value == "=AF5-AH5-AI5"
    assert item["AN6"].value == "=AG6*L6"
    assert item["I7"].value == "合计"
    assert item["L7"].value == "=SUBTOTAL(9,L4:L6)"
    assert item["A5"].fill.fgColor.rgb == item["A4"].fill.fgColor.rgb
    assert item["C5"].font.color.type == "rgb"
    assert item["C5"].font.color.rgb.endswith("0000FF")

    review = rendered["正单评审表 "]
    assert review["C6"].value == "='Iteam表'!B5"
    assert review["E7"].value == "='Iteam表'!E6"
    assert review["S7"].value == "=I7*R7"
    assert review["F8"].value == "=Iteam表!I7"
    order = rendered["接单表"]
    assert order["A6"].value == "='Iteam表'!B5"
    assert order["C7"].value == "='Iteam表'!E6"
    assert order["J7"].value == "=H7/I7"
    assert order["E8"].value == "=Iteam表!I7"
    assert rendered["历史资料"]["A1"].value == "必须保留"
    rendered.close()


def test_complete_append_skips_group_summary_rows_when_selecting_style_and_boundary(
    tmp_path: Path,
):
    source_path = tmp_path / "分组排期.xlsx"
    output_path = tmp_path / "分组排期_新单.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "接单表"
    worksheet.append(["PO号", "货号", "数量", "箱数"])
    worksheet.append(["PO-OLD", "ITEM-OLD", 12, "=C2/4"])
    worksheet["A2"].fill = PatternFill("solid", fgColor="DDEBF7")
    worksheet.merge_cells("A3:D3")
    worksheet["A3"] = "分组汇总"
    worksheet["A3"].fill = PatternFill("solid", fgColor="FFF2CC")
    worksheet["A4"] = "合计"
    worksheet["C4"] = "=SUBTOTAL(9,C2:C2)"
    workbook.save(source_path)
    workbook.close()

    new_order_excel.append_records_to_workbook(
        source_path,
        output_path,
        [{"po_no": "PO-NEW", "item_no": "ITEM-NEW", "quantity": 20}],
        {
            "po_no": ("PO号",),
            "item_no": ("货号",),
            "quantity": ("数量",),
            "cartons": ("箱数",),
        },
        filename=source_path.name,
        sheet_names=("接单表",),
    )

    rendered = load_workbook(output_path, data_only=False)
    worksheet = rendered["接单表"]
    assert worksheet["A2"].value == "PO-OLD"
    assert worksheet["A3"].value == "PO-NEW"
    assert worksheet["D3"].value == "=C3/4"
    assert worksheet["A3"].fill.fgColor.rgb == worksheet["A2"].fill.fgColor.rgb
    assert worksheet["A4"].value == "分组汇总"
    assert worksheet["A5"].value == "合计"
    assert worksheet["C5"].value == "=SUBTOTAL(9,C2:C3)"
    rendered.close()


def test_complete_append_keeps_formula_with_record_fallback(tmp_path: Path):
    source_path = tmp_path / "公式兜底排期.xlsx"
    output_path = tmp_path / "公式兜底排期_新单.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "接单表"
    worksheet.append(["数量", "装箱", "箱数"])
    worksheet.append([12, 4, "=A2/B2"])
    workbook.save(source_path)
    workbook.close()

    new_order_excel.append_records_to_workbook(
        source_path,
        output_path,
        [{"quantity": 60, "pack_qty": "3P/12", "cartons": 5}],
        {
            "quantity": ("数量",),
            "pack_qty": ("装箱",),
            "cartons": ("箱数",),
        },
        filename=source_path.name,
        sheet_names=("接单表",),
        formula_fallback_fields=("cartons",),
    )

    rendered = load_workbook(output_path, data_only=False)
    assert rendered["接单表"]["C3"].value == "=IFERROR(A3/B3,5)"
    rendered.close()


def test_seasons_composite_pack_uses_outer_carton_quantity():
    assert shixin_schedule.outer_pack_number("4P/16") == 16
    assert shixin_schedule.outer_pack_number("3P/12") == 12
    assert shixin_schedule.outer_pack_number("24/144") == 144


def test_seasons_appends_to_reserved_rows_before_first_total_and_stays_in_first_section(
    tmp_path: Path,
):
    source_path = tmp_path / "施信排期.xlsx"
    output_path = tmp_path / "施信排期_新单.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "正单评审表"
    worksheet.append(["2026年施信产品评审表"])
    worksheet.append([
        "证书", "客出单日期", "预备单号（OQF NO）", "O/C NO", "PO.NO",
        "客名/國家", "產品編號", "產品名称", "數量", "装箱", "箱数",
        "行Q", "客Q", "Q货情况", "客要求走货期", "包装要求", "国家标准",
        "备注", "生产车间", "上系统", "单价", "金额HKD", "客货号",
    ])
    worksheet.append([None] * 23)
    worksheet.append([
        "证书A", date(2026, 5, 16), "QF-OLD", "OC-OLD", "PO-OLD", "BIG W",
        "ITEM-OLD", "历史产品", 150, 6, "=I4/J4", None, None, "自检ok",
        date(2026, 7, 1), "客卡", "美国标准", "历史备注", "历史车间", "已上",
        31.34, "=I4*U4", "客户货号",
    ])
    worksheet["A4"].fill = PatternFill("solid", fgColor="DDEBF7")
    worksheet["A4"].font = Font(name="宋体", color="0000FF")
    worksheet.append([None] * 23)
    worksheet.append([None] * 20 + ["合计：", "=SUM(V3:V5)", "=V6/10000"])
    worksheet.append([None] * 20 + ["接单总计：", "=V6+V12", "=V7/10000"])
    worksheet.append([None] * 23)
    worksheet.append([None] * 23)
    worksheet.append([None] * 23)
    worksheet.append([None] * 23)
    worksheet.append([
        None, date(2025, 12, 15), "QF-LATER", "OC-LATER", "PO-LATER", "WORLD MARKET",
        "ITEM-LATER", "后段走货订单", 1002, 2, 501,
    ])
    workbook.save(source_path)
    workbook.close()

    aliases: dict[str, list[str]] = {}
    for label, field in shixin_schedule.FIELDS.items():
        aliases.setdefault(field, []).append(label)
    for field, label in shixin_schedule.EXPORT_FIELDS:
        aliases.setdefault(field, []).append(label)

    result = new_order_excel.append_records_to_workbook(
        source_path,
        output_path,
        [
            {
                "order_date": "2026-08-12",
                "oqf_no": "QF-NEW-1",
                "oc_no": "OC-NEW-1",
                "po_no": "PO-NEW-1",
                "customer": "SEASONS USA",
                "item_no": "ITEM-NEW-1",
                "product_name": "新产品1",
                "quantity": 60,
                "pack_qty": 6,
                "cartons": 10,
                "ship_date": "2026-09-30",
                "unit_price": 12.5,
                "amount": 750,
            },
            {
                "order_date": "2026-08-12",
                "oqf_no": "QF-NEW-2",
                "oc_no": "OC-NEW-2",
                "po_no": "PO-NEW-2",
                "customer": "SEASONS USA",
                "item_no": "ITEM-NEW-2",
                "product_name": "新产品2",
                "quantity": 48,
                "pack_qty": 6,
                "cartons": 8,
                "ship_date": "2026-10-15",
                "unit_price": 10,
                "amount": 480,
            },
        ],
        aliases,
        filename=source_path.name,
        sheet_names=shixin_schedule.SEASONS_SHEETS,
        formula_fallback_fields=("cartons",),
        first_total_section=True,
    )

    rendered = load_workbook(output_path, data_only=False)
    worksheet = rendered["正单评审表"]
    assert result["header_row"] == 2
    assert result["insert_row"] == 5
    assert result["section_total_row"] == 6
    assert result["reused_blank_rows"] == 1
    assert worksheet["C5"].value == "QF-NEW-1"
    assert worksheet["E5"].value == "PO-NEW-1"
    assert worksheet["K5"].value == "=IFERROR(I5/J5,10)"
    assert worksheet["V5"].value == "=I5*U5"
    assert worksheet["C6"].value == "QF-NEW-2"
    assert worksheet["E6"].value == "PO-NEW-2"
    assert worksheet["K6"].value == "=IFERROR(I6/J6,8)"
    assert worksheet["V6"].value == "=I6*U6"
    assert worksheet["A5"].fill.fgColor.rgb == worksheet["A4"].fill.fgColor.rgb
    assert worksheet["A5"].font.name == worksheet["A4"].font.name
    assert worksheet["A6"].fill.fgColor.rgb == worksheet["A4"].fill.fgColor.rgb
    assert worksheet["U7"].value == "合计："
    assert worksheet["V7"].value == "=SUM(V3:V6)"
    assert worksheet["V8"].value == "=V7+V13"
    assert worksheet["E13"].value == "PO-LATER"
    rendered.close()


def test_seasons_export_synchronizes_review_and_matching_blow_mold_lane(tmp_path: Path):
    workbook = Workbook()
    review = workbook.active
    review.title = "正单评审表"
    review_headers = [
        "证书", "客出单日期", "预备单号（OQF NO）", "O/C NO", "PO.NO",
        "客名/國家", "產品編號", "產品名称", "數量", "装箱", "箱数",
        "行Q", "客Q", "Q货情况", "客要求走货期", "包装要求", "国家标准",
        "备注", "生产车间", "上系统", "单价", "金额HKD",
    ]
    for column, value in enumerate(review_headers, 1):
        review.cell(2, column, value)
    review_row = [
        None, datetime(2026, 1, 15), "QF-OLD", "ZE-OLD", "PO-OLD",
        "ALBERTSONS", "W85340", "四层橙色南瓜堆", 50, 1, "=I4/J4",
        None, None, None, datetime(2026, 4, 22), "客彩贴", "美国标准",
        "不包电", None, "已上", 288.32, "=I4*U4",
    ]
    for column, value in enumerate(review_row, 1):
        review.cell(4, column, value)
    review["U6"] = "合计："
    review["V6"] = "=SUM(V3:V5)"

    order = workbook.create_sheet("吹气、PU接单表")
    order_headers = [
        None, "客出单日期", "预备单号（OQF NO）", "O/C NO", "PO.NO",
        "客名/國家", "產品編號", "產品名称", "數量", "装箱", "单价", "金额",
        "包装", "備註", "生产车间", "客要求走货期",
    ]
    for column, value in enumerate(order_headers, 1):
        order.cell(2, column, value)
    for column, value in enumerate([
        None, datetime(2026, 1, 21), "QF-OLD", "ZE-OLD", "PO-OLD",
        "ALBERTSONS", "W85340", "四层橙色南瓜堆", 50, 1, 288.32,
        "=I4*K4", "客彩贴", "不包电", None, datetime(2026, 4, 22),
    ], 1):
        order.cell(4, column, value)
    order["K5"] = "合计"
    order["L5"] = "=SUM(L4:L4)"

    item = workbook.create_sheet("吹气系列ITEM表")
    item_headers = [
        "证书", "客出单日期", "预备单号（OQF NO）", "O/C NO", "PO.NO",
        "客名/國家", "產品編號", "產品名稱", "數量", "包装", "備註",
        "生产车间", "客要求走货期",
    ]
    for column, value in enumerate(item_headers, 1):
        item.cell(3, column, value)
    for column, value in enumerate([
        None, datetime(2026, 1, 21), "QF-OLD", "ZE-OLD", "PO-OLD",
        "ALBERTSONS", "W85340", "四层橙色南瓜堆", 50, "客彩贴", "不包电",
        None, datetime(2026, 4, 22),
    ], 1):
        item.cell(4, column, value)
    item["H5"] = "合计"
    item["I5"] = "=SUM(I4:I4)"

    source = BytesIO()
    workbook.save(source)
    workbook.close()
    output_path = tmp_path / "SEASONS排期_新单.xlsx"
    record = {
        "order_date": "2026-08-14",
        "oqf_no": "QF16098253",
        "oc_no": "ZE618957",
        "po_no": "ZPO1621463",
        "customer": "ALBERTSONS",
        "item_no": "W85340",
        "product_name": "四层橙色南瓜堆",
        "quantity": 50,
        "pack_qty": 1,
        "cartons": 50,
        "ship_date": "2026-04-22",
        "packaging": "客彩贴",
        "standard": "美国标准",
        "notes": "不包电",
        "unit_price": 288.32,
        "amount": 14416,
    }

    _export_prepared(
        customer_code="seasons",
        prepared=PreparedBatch([record], [], "正单评审表"),
        schedule_file_name="SEASONS排期.xlsx",
        schedule_content=source.getvalue(),
        output_path=output_path,
    )

    rendered = load_workbook(output_path, data_only=False)
    assert rendered["正单评审表"]["E5"].value == "ZPO1621463"
    assert rendered["正单评审表"]["V5"].value == "=I5*U5"
    assert rendered["吹气、PU接单表"]["E5"].value == "ZPO1621463"
    assert rendered["吹气、PU接单表"]["L5"].value == "=I5*K5"
    assert rendered["吹气系列ITEM表"]["E5"].value == "ZPO1621463"
    assert rendered["吹气系列ITEM表"]["G5"].value == "W85340"
    rendered.close()


def test_yinhui_identifiers_are_written_as_text(tmp_path: Path):
    source_path = tmp_path / "银辉排期.xlsx"
    output_path = tmp_path / "银辉排期_新单.xlsx"
    workbook = _make_yinhui_three_sheet_template()
    workbook.save(source_path)
    workbook.close()

    yinhui_schedule.create_export(
        [{
            "so_no": "2300000463-780",
            "contract_no": "4500002406",
            "item_no": "5122208863801",
            "quantity": 20,
        }],
        output_path,
        source_path,
        template_filename=source_path.name,
    )

    rendered = load_workbook(output_path, data_only=False)
    worksheet = rendered["Iteam表"]
    for coordinate in ("B6", "C6", "E6"):
        assert worksheet[coordinate].data_type == "s"
        assert worksheet[coordinate].number_format == worksheet[f"{coordinate[0]}5"].number_format
    assert worksheet["I6"].value == '=IFERROR(G6/H6,"")'
    rendered.close()


def test_seasons_schedule_families_are_kept_in_separate_lanes():
    assert shixin_schedule.schedule_family_from_sheets(["正单评审表"]) == "seasons"
    assert shixin_schedule.schedule_family_from_sheets(["手掌", "面具"]) == "internal"


def test_maxx_and_shushupapa_apply_their_seven_day_dates():
    maxx = multi_schedule._line_values(
        {"client": "maxx", "po_number": "MAXX-1", "ship_date": "2026-08-20"},
        {"item_code": "100", "description": "Maxx item", "quantity": 10},
    )
    shushupapa = multi_schedule._line_values(
        {"client": "shushupapa", "po_number": "SSP-1", "ship_date": "2026-08-20"},
        {"item_code": "200", "description": "Shushu item", "quantity": 10},
    )

    assert maxx["complete_date"] == "2026-08-13"
    assert maxx["inspection_date"] == "2026-08-13"
    assert shushupapa["complete_date"] is None
    assert shushupapa["inspection_date"] == "2026-08-13"


def test_multi_customer_revision_dedup_keeps_highest_revision():
    selected, report = _dedupe_multi_orders([
        {"po_number": "PO-1", "filename": "PO-1.pdf"},
        {"po_number": "PO-1", "filename": "PO-1_REV2.pdf"},
    ])

    assert [item["filename"] for item in selected] == ["PO-1_REV2.pdf"]
    assert "PO-1.pdf" in report[0]


def test_common_preview_contract_maps_each_customer_and_blocks_high_risk_flags():
    assert _record_fields("edu", {"customer_po": "E-1", "item_no": "A"})["po_no"] == "E-1"
    assert _record_fields("360", {"production_no": "RL-1", "item_full": "B"})["contract_no"] == "RL-1"
    assert _record_fields("yinhui", {"contract_no": "Y-1"})["po_no"] == "Y-1"
    assert _record_fields("seasons", {"oqf_no": "QF-1"})["po_no"] == "QF-1"
    assert _record_fields("maxx", {"po_number": "M-1"})["po_no"] == "M-1"
    assert _record_fields("shushupapa", {"po_number": "S-1"})["po_no"] == "S-1"

    issues = _issues(
        {"flags": [{"level": "high", "code": "missing_item", "text": "缺货号"}]},
        "maxx-1",
    )
    assert issues[0]["severity"] == "blocked"
    assert issues[0]["can_skip"] is False


def test_existing_orders_and_data_risks_support_controlled_manual_resolution():
    issues = _issues(
        {
            "po_number": "PO-EXISTING",
            "_duplicate_existing": True,
            "flags": [{"level": "high", "code": "missing_item", "text": "缺货号"}],
        },
        "maxx-1",
    )
    duplicate = next(issue for issue in issues if issue["code"] == "duplicate_existing_order")
    hard_blocker = next(issue for issue in issues if issue["code"] == "missing_item")

    assert duplicate["severity"] == "blocked"
    assert duplicate["can_skip"] is True
    assert duplicate["skip_label"] == "测试阶段确认重复导入当前或历史排期已有订单"
    assert hard_blocker["can_skip"] is False

    preview = {"rows": [{"issues": issues}]}
    with pytest.raises(HuaxingCustomerOrderError, match="缺货号"):
        _validate_skips(preview, {duplicate["skip_key"]})
    _validate_skips({"rows": [{"issues": [duplicate]}]}, {duplicate["skip_key"]})
    _validate_skips(
        preview,
        {duplicate["skip_key"], hard_blocker["skip_key"]},
    )
