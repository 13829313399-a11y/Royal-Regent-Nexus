from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.services.injection_scheduling_excel import parse_injection_scheduling_workbook
from app.services.injection_scheduling_profiles import (
    BUILTIN_IMPORT_PROFILES,
    GROUP_UNIFIED_PLAN_HEADERS,
)
from app.services.injection_scheduling_unified_template import (
    TEMPLATE_RESOURCE,
    TEMPLATE_SHA256,
    TEMPLATE_VERSION,
    unified_template_download,
)


def _fixture(*, changed_header: bool = False, multiple_running: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "计划导入"
    sheet["A1"] = "Royal Regent Nexus｜集团统一注塑排产导入模板"
    sheet["A2"] = "模板版本"
    sheet["B2"] = TEMPLATE_VERSION
    for index, header in enumerate(GROUP_UNIFIED_PLAN_HEADERS, start=1):
        sheet.cell(5, index, header)
    if changed_header:
        sheet["E5"] = "产品"

    for row in range(6, 15):
        sheet[f"A{row}"] = f'=IF(B{row}="","",ROW()-5)'
        sheet[f"I{row}"] = f'=IF(G{row}="","",MAX(G{row}-IF(H{row}="",0,H{row}),0))'
        sheet[f"AD{row}"] = f'=IF(B{row}="","",IF(G{row}>0,"可导入","订单数量错误"))'

    backlog = {
        "B": "000123",
        "C": "000045",
        "D": "MOLD-01",
        "E": "浅蓝外壳",
        "F": "007",
        "G": 100,
        "J": datetime(2026, 9, 1),  # noqa: DTZ001 - Excel local business date
        "K": "特急",
        "L": 7,
        "M": 85,
        "N": 102,
        "O": 50,
        "P": "ABS",
        "Q": "浅蓝",
        "R": "浅",
        "S": "双臂",
        "T": "气剪",
        "U": "透明料|自定义工艺",
        "Y": "待排",
        "AB": "否",
    }
    for column, value in backlog.items():
        sheet[f"{column}6"] = value

    scheduled = {
        "B": "SO-002",
        "D": "MOLD-02",
        "E": "黑色底座",
        "G": 200,
        "H": 20,
        "J": datetime(2026, 9, 3),  # noqa: DTZ001 - Excel local business date
        "K": "急单",
        "X": "A-12",
        "Y": "生产中",
        "AB": "是",
    }
    for column, value in scheduled.items():
        sheet[f"{column}7"] = value

    if multiple_running:
        for column, value in {
            "B": "SO-003",
            "D": "MOLD-03",
            "E": "后续任务",
            "G": 80,
            "J": datetime(2026, 9, 4),  # noqa: DTZ001
            "X": "A-12",
            "Y": "生产中",
            "AB": "是",
        }.items():
            sheet[f"{column}8"] = value

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_shipped_template_contract_and_download_identity():
    assert TEMPLATE_RESOURCE.is_file()
    assert __import__("hashlib").sha256(TEMPLATE_RESOURCE.read_bytes()).hexdigest() == TEMPLATE_SHA256
    workbook = load_workbook(TEMPLATE_RESOURCE, read_only=True, data_only=False)
    assert workbook.sheetnames == ["计划导入", "填写说明", "字段说明", "填写示例", "枚举字典"]
    sheet = workbook["计划导入"]
    assert sheet["B2"].value == TEMPLATE_VERSION
    assert tuple(sheet.cell(5, index).value for index in range(1, 31)) == GROUP_UNIFIED_PLAN_HEADERS
    path, filename = unified_template_download("huakang-b")
    assert path == TEMPLATE_RESOURCE
    assert filename == "华康B_统一注塑排产导入模板_RR-ISP-1.0.xlsx"


def test_empty_shipped_template_has_no_formula_only_business_rows():
    normalized, issues = parse_injection_scheduling_workbook(
        TEMPLATE_RESOURCE.read_bytes(),
        TEMPLATE_RESOURCE.name,
        factory_id="huaxing",
        profiles=tuple(
            profile
            for profile in BUILTIN_IMPORT_PROFILES
            if profile.profile_code == "group_unified_plan_v1"
        ),
    )

    assert normalized["summary"]["backlog_count"] == 0
    assert normalized["summary"]["scheduled_baseline_count"] == 0
    assert normalized["summary"]["can_confirm"] is False
    assert any(issue["code"] == "NO_IMPORT_ROWS" for issue in issues)


def test_unified_parser_preserves_identifiers_and_splits_backlog_from_baseline():
    normalized, issues = parse_injection_scheduling_workbook(
        _fixture(),
        "集团统一注塑排产导入模板.xlsx",
        factory_id="huakang-b",
        system_machine_codes={"A-12"},
        profiles=tuple(
            profile
            for profile in BUILTIN_IMPORT_PROFILES
            if profile.profile_code == "group_unified_plan_v1"
        ),
    )
    assert normalized["profile"]["profile_code"] == "group_unified_plan_v1"
    assert normalized["summary"]["backlog_count"] == 1
    assert normalized["summary"]["scheduled_baseline_count"] == 1
    assert normalized["summary"]["invalid_row_count"] == 0
    backlog = normalized["backlog_orders"][0]
    assert backlog["order_no"] == "000123"
    assert backlog["item_no"] == "000045"
    assert backlog["warehouse_text"] == "007"
    assert backlog["completed_quantity"] == 0
    assert backlog["priority_code"] == "CRITICAL"
    assert backlog["execution_status"] == "BACKLOG"
    scheduled = normalized["scheduled_baseline_tasks"][0]
    assert scheduled["machine_code"] == "A-12"
    assert scheduled["execution_status"] == "RUNNING"
    assert scheduled["locked"] is True
    assert scheduled["planned_start"] == ""
    assert scheduled["planned_finish"] == ""
    assert any(issue["code"] == "UNKNOWN_PROCESS_TAG" for issue in issues)
    assert not any(issue["blocking"] for issue in issues)


def test_unified_parser_rejects_modified_headers_and_multiple_running_tasks():
    normalized, issues = parse_injection_scheduling_workbook(
        _fixture(changed_header=True),
        "modified.xlsx",
        factory_id="huakang-b",
        system_machine_codes={"A-12"},
        profiles=tuple(
            profile
            for profile in BUILTIN_IMPORT_PROFILES
            if profile.profile_code == "group_unified_plan_v1"
        ),
    )
    assert normalized["summary"]["can_confirm"] is False
    assert any(issue["code"] == "UNIFIED_TEMPLATE_STRUCTURE_MODIFIED" for issue in issues)

    normalized, issues = parse_injection_scheduling_workbook(
        _fixture(multiple_running=True),
        "running.xlsx",
        factory_id="huakang-b",
        system_machine_codes={"A-12"},
        profiles=tuple(
            profile
            for profile in BUILTIN_IMPORT_PROFILES
            if profile.profile_code == "group_unified_plan_v1"
        ),
    )
    assert normalized["summary"]["can_confirm"] is False
    assert any(issue["code"] == "MULTIPLE_RUNNING_TASKS" for issue in issues)


def test_unified_repeated_order_rows_share_order_identity_but_keep_split_identity():
    workbook = load_workbook(BytesIO(_fixture()))
    sheet = workbook["计划导入"]
    for column in range(2, 30):
        sheet.cell(8, column, sheet.cell(6, column).value)
    output = BytesIO()
    workbook.save(output)

    normalized, issues = parse_injection_scheduling_workbook(
        output.getvalue(),
        "repeated-order.xlsx",
        factory_id="huakang-b",
        system_machine_codes={"A-12"},
        profiles=tuple(
            profile
            for profile in BUILTIN_IMPORT_PROFILES
            if profile.profile_code == "group_unified_plan_v1"
        ),
    )

    repeated = [
        row for row in normalized["backlog_orders"] if row["order_no"] == "000123"
    ]
    assert len(repeated) == 2
    assert repeated[0]["stable_order_key"] == repeated[1]["stable_order_key"]
    assert repeated[0]["stable_row_key"] != repeated[1]["stable_row_key"]
    assert not any(issue["blocking"] for issue in issues)
