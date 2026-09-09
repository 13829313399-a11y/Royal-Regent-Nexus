"""Real generated native files, not mocked extraction; twenty diverse samples."""
from __future__ import annotations

import datetime as dt
import hashlib
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName

from app.services.document_tools import office_engine as office
from app.services.document_tools.document_ir import Block, Cell, DocumentIR, Page, SourceAnchor, Table, Cancelled, ToolError


CASES = ["word-basic", "word-horizontal", "word-vertical", "word-rectangle", "word-nested",
         "word-header", "word-long", "word-notes", "word-multi", "word-identifiers",
         "excel-basic", "excel-merged", "excel-hidden", "excel-formula", "excel-dates",
         "excel-formats", "excel-print", "excel-named", "excel-charts", "excel-wide"]


def generate_sample(folder: Path, case: str) -> tuple[Path, dict]:
    folder.mkdir(parents=True, exist_ok=True)
    options = {}
    if case.startswith("word"):
        path = folder / (case + " 中文 (样本).docx")
        doc = Document()
        doc.add_heading("华康原生办公文档测试", 0)
        doc.add_paragraph("原文说明：本文件为自建测试样本，不含真实业务资料。数量 12.50，编号 0000123。")
        table = doc.add_table(rows=5, cols=4)
        table.style = "Table Grid"
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                cell.text = ["编号", "产品", "数量", "说明"][c] if r == 0 else f"R{r}C{c} 文本"
        table.cell(1, 0).text = "00001234567890123456789"
        table.cell(1, 2).text = "12.50"
        header = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header)
        if case == "word-horizontal":
            table.cell(2, 0).merge(table.cell(2, 2)).text = "横向合并"
        if case == "word-vertical":
            table.cell(1, 1).merge(table.cell(3, 1)).text = "纵向合并"
        if case == "word-rectangle":
            table.cell(1, 1).merge(table.cell(3, 2)).text = "矩形合并"
        if case == "word-nested":
            nested = table.cell(2, 3).add_table(rows=2, cols=2)
            nested.cell(0, 0).text = "嵌套标题"
            nested.cell(1, 1).text = "嵌套值 2468"
        if case == "word-header":
            options["include_headers_footers"] = True
            header_table = doc.sections[0].header.add_table(rows=1, cols=2, width=Inches(5))
            header_table.cell(0, 0).text = "页眉表格"
            doc.sections[0].footer.paragraphs[0].text = "页脚说明"
        if case == "word-long":
            for n in range(55):
                doc.add_paragraph(f"长文段落 {n+1}：保留阅读顺序与正常换行。" * 3)
            doc.add_page_break()
            doc.add_paragraph("明确分页之后的结束语")
        if case == "word-notes":
            for n in range(8):
                doc.add_paragraph(f"注意事项 {n+1}：温度 23.5℃，误差 ±0.2 mm。")
            options["word_mode"] = "structure"
        if case == "word-multi":
            doc.add_heading("第二张业务表", 1)
            second = doc.add_table(rows=3, cols=2)
            second.cell(0, 0).text = "不同语义表头"
            second.cell(2, 1).text = "不能仅凭列数自动合并"
        if case == "word-identifiers":
            table.cell(2, 0).text = "=HYPERLINK(\"https://invalid.example\")"
            table.cell(2, 2).text = "1,234"
            table.cell(3, 2).text = "1.234"
        doc.save(path)
    else:
        path = folder / (case + " 中文 (样本).xlsx")
        book = Workbook()
        sheet = book.active
        sheet.title = "订单明细"
        sheet.append(["编号", "品名", "数量", "单价", "备注"])
        for n in range(1, 11):
            sheet.append([f"{n:05}", f"产品 {n}", n*2, 12.5, "中文说明，单位 mm"])
        sheet["A2"] = "00001234567890123456789"
        sheet.column_dimensions["A"].width = 28
        sheet["D2"].number_format = "#,##0.00"
        sheet["B2"].fill = PatternFill("solid", fgColor="DDEEFF")
        if case == "excel-merged":
            sheet.merge_cells("B3:C5")
            sheet["B3"] = "合并范围"
        if case == "excel-hidden":
            sheet.row_dimensions[3].hidden = True
            sheet.column_dimensions["D"].hidden = True
            hidden = book.create_sheet("隐藏工作表")
            hidden["A1"] = "不得默认导出"
            hidden.sheet_state = "hidden"
        if case == "excel-formula":
            sheet["C2"] = "=SUM(C3:C5)"
            sheet["D3"] = "=_xlfn.UNKNOWN(1)"
        if case == "excel-dates":
            sheet["C2"] = dt.datetime(2026, 9, 8, 10, 25)
            sheet["C2"].number_format = "yyyy-mm-dd hh:mm:ss"
            sheet["C3"] = dt.date(2026, 12, 31)
            sheet["C3"].number_format = "yyyy-mm-dd"
        if case == "excel-formats":
            sheet["C2"], sheet["C2"].number_format = 0.125, "0.00%"
            sheet["D3"], sheet["D3"].number_format = -1234.5, '#,##0.00;(#,##0.00)'
            sheet["C4"], sheet["C4"].number_format = 123, "000000"
        if case == "excel-print":
            sheet.print_area = "A1:D6"
            sheet.print_title_rows = "1:1"
            sheet.page_setup.orientation = "landscape"
        if case == "excel-named":
            book.defined_names.add(DefinedName("所选区域", attr_text="'订单明细'!$A$1:$C$5"))
            options["range"] = "所选区域"
        if case == "excel-charts":
            chart = BarChart()
            chart.add_data(Reference(sheet, min_col=3, min_row=1, max_row=11), titles_from_data=True)
            sheet.add_chart(chart, "G2")
        if case == "excel-wide":
            for c in range(6, 18):
                sheet.cell(1, c, f"列 {c}")
                for r in range(2, 15):
                    sheet.cell(r, c, f"宽表 R{r}C{c}")
        book.save(path)
    return path, options


def generate_corpus(folder: Path):
    return [generate_sample(folder, case) for case in CASES]


@pytest.mark.parametrize("case", CASES)
def test_twenty_real_native_samples(tmp_path, case):
    source, options = generate_sample(tmp_path, case)
    original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    progress = []
    _, ir = office._extract(source, options, tmp_path, lambda *args: progress.append(args), lambda: False)
    assert ir.tables and progress
    assert any(c.display_text == "00001234567890123456789" for t in ir.tables for c in t.cells)
    assert len({c.id for t in ir.tables for c in t.cells}) == sum(len(t.cells) for t in ir.tables)
    if case == "excel-basic":
        assert next(c for c in ir.tables[0].cells if c.source.cell == "A1").style["background"] is None
    if case.startswith("word"):
        output = office.write_xlsx(ir, tmp_path / "result.xlsx", options)
        result = load_workbook(output)
        assert len(result.worksheets) >= len(ir.tables)
        assert result.worksheets[0]["A2"].value == "00001234567890123456789"
        if case == "word-horizontal":
            assert "A3:C3" in result.worksheets[0].merged_cells
        if case == "word-vertical":
            assert "B2:B4" in result.worksheets[0].merged_cells
        if case == "word-rectangle":
            assert "B2:C4" in result.worksheets[0].merged_cells
        if case == "word-nested":
            assert len(ir.tables) == 2
            assert ir.tables[1].parent_cell_id
            assert any(c.display_text == "嵌套值 2468" for c in ir.tables[1].cells)
        if case == "word-header":
            assert any("header" in t.source.block_id for t in ir.tables)
        if case == "word-identifiers":
            assert result.worksheets[0]["A3"].data_type == "s"
        result.close()
    else:
        output = office.write_docx(ir, tmp_path / "result.docx", options)
        result = Document(output)
        assert result.tables
        assert any("00001234567890123456789" in cell.text for table in result.tables for row in table.rows for cell in row.cells)
        if case == "excel-hidden":
            assert len(ir.tables) == 1 and ir.tables[0].column_count == 4 and ir.tables[0].row_count == 10
        if case == "excel-formula":
            assert any(i.code == "FORMULA_CACHE_MISSING" for i in ir.issues)
            assert any(c.resolution == "unresolved" for t in ir.tables for c in t.cells)
        if case == "excel-print":
            assert (ir.tables[0].row_count, ir.tables[0].column_count) == (6, 4)
        if case == "excel-named":
            assert (ir.tables[0].row_count, ir.tables[0].column_count) == (5, 3)
        if case == "excel-formats":
            values = {c.source.cell: c.display_text for c in ir.tables[0].cells}
            assert values["C2"] == "12.50%"
            assert values["C4"] == "000123"
            assert values["D3"] == "(1,234.50)"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == original_hash
    assert zipfile.is_zipfile(output)
    assert ir.mappings


def test_cancelled_before_extract(tmp_path):
    source, options = generate_sample(tmp_path, "word-basic")
    with pytest.raises(Cancelled):
        office._extract(source, options, tmp_path, lambda *args: None, lambda: True)


def test_bad_file_is_actionable(tmp_path):
    source = tmp_path / "broken.docx"
    source.write_bytes(b"broken")
    with pytest.raises(ToolError, match="Office"):
        office._extract(source, {}, tmp_path, lambda *args: None, lambda: False)


def test_formula_text_explicit(tmp_path):
    source, _ = generate_sample(tmp_path, "excel-formula")
    _, ir = office._extract(source, {"formula_mode": "formula"}, tmp_path, lambda *args: None, lambda: False)
    cell = next(c for c in ir.tables[0].cells if c.source.cell == "C2")
    assert cell.display_text == "=SUM(C3:C5)"
    assert cell.resolution == "resolved"


def test_hidden_include_explicit(tmp_path):
    source, _ = generate_sample(tmp_path, "excel-hidden")
    _, ir = office._extract(source, {"include_hidden": True}, tmp_path, lambda *args: None, lambda: False)
    assert len(ir.tables) == 2
    assert ir.tables[0].column_count == 5


def test_missing_engine_keeps_native_inspection(tmp_path, monkeypatch):
    source, _ = generate_sample(tmp_path, "word-basic")
    monkeypatch.setattr(office, "office_executable", lambda: None)
    result = office.inspect_office(source, {}, tmp_path / "job", lambda *args: None, lambda: False)
    assert result.ir.tables
    assert any(i.code == "OFFICE_UNAVAILABLE" for i in result.ir.issues)


def test_manual_revision_not_overwritten_by_recalculation(tmp_path, monkeypatch):
    source, _ = generate_sample(tmp_path, "excel-formula")
    _, ir = office._extract(source, {}, tmp_path, lambda *args: None, lambda: False)
    cell = next(c for c in ir.tables[0].cells if c.source.cell == "C2")
    cell.value, cell.display_text, cell.resolution, cell.source.method = "42", "42", "manually_confirmed", "manual"
    monkeypatch.setattr(office, "office_executable", lambda: None)
    result = office.convert_office(source, "excel_to_word", {}, tmp_path / "job", lambda *args: None, lambda: False, ir=ir)
    assert next(c for c in result.ir.tables[0].cells if c.source.cell == "C2").display_text == "42"


@pytest.mark.parametrize("case", ["word-basic", "word-nested", "word-header", "excel-charts"])
def test_pdf_revision_patches_only_native_xml(tmp_path, case):
    source, options = generate_sample(tmp_path, case)
    _, ir = office._extract(source, options, tmp_path, lambda *args: None, lambda: False)
    table = ir.tables[-1] if case in ("word-nested", "word-header") else ir.tables[0]
    cell = next(c for c in table.cells if c.display_text)
    old_text = cell.display_text
    cell.display_text, cell.value = "人工核验 000099", "人工核验 000099"
    cell.source.method, cell.resolution = "manual", "manually_confirmed"
    output = office._corrected_copy(source, ir, tmp_path / ("corrected" + source.suffix))
    _, fresh = office._extract(output, options, tmp_path, lambda *args: None, lambda: False)
    assert any(c.display_text == "人工核验 000099" for t in fresh.tables for c in t.cells)
    assert not any(c.display_text == old_text and c.source.block_id == cell.source.block_id and c.source.cell == cell.source.cell for t in fresh.tables for c in t.cells)
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output) as revised:
        assert set(original.namelist()) == set(revised.namelist())
        changed = {name for name in original.namelist() if original.read(name) != revised.read(name)}
        if source.suffix == ".xlsx":
            assert changed <= {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}
            assert original.read("xl/charts/chart1.xml") == revised.read("xl/charts/chart1.xml")
        else:
            assert len(changed) == 1


def test_calc_display_preserves_manual_and_reports_changed_cache(tmp_path):
    source, options = generate_sample(tmp_path, "excel-formula")
    _, ir = office._extract(source, options, tmp_path, lambda *args: None, lambda: False)
    formula = next(c for c in ir.tables[0].cells if c.source.cell == "C2")
    metadata = {"display_cells": {"订单明细": {"2:3": {"display": "24", "error": 0}}}}
    office._apply_calc_metadata(ir, metadata)
    assert formula.display_text == "24" and formula.resolution == "resolved"
    assert not any(i.target_id == formula.id and i.code == "FORMULA_CACHE_MISSING" for i in ir.issues)
    formula.source.method = "manual"
    metadata = {"display_cells": {"订单明细": {"2:3": {"display": "999", "error": 0}}}}
    office._apply_calc_metadata(ir, metadata)
    assert formula.display_text == "24"


def make_layout_sample(folder):
    ir = DocumentIR(source_type="pdf", pages=[Page(page_index=n, display_page_number=n+1, width_pt=595, height_pt=842) for n in range(2)])
    ir.blocks = [Block(id=f"page{n}-{c}", text=f"PAGE {n+1} COLUMN {c+1} CONTENT 000123", source=SourceAnchor(page_index=n, bbox_pt=[40+c*280, 80, 285+c*280, 110]), style={"font_size": 11}) for n in range(2) for c in range(2)]
    ir.tables.append(Table(id="positioned-table", row_count=2, column_count=2,
                           source=SourceAnchor(page_index=0, bbox_pt=[40, 150, 450, 210]),
                           cells=[Cell(id=f"p-cell{r}{c}", row=r, column=c, display_text=f"TABLE R{r}C{c}", value=f"TABLE R{r}C{c}") for r in range(2) for c in range(2)]))
    ir.blocks.append(Block(id="table-block", kind="table", table_id="positioned-table", source=ir.tables[0].source))
    return office.write_docx(ir, folder / "positioned.docx", {"layout_mode": "layout"})


def test_layout_has_real_textboxes_and_table(tmp_path):
    path = make_layout_sample(tmp_path)
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode()
        assert xml.count("<w:txbxContent>") == 5
        assert "margin-left:320.000pt" in xml
        assert "TABLE R1C1" in xml and "<w:tbl>" in xml
        assert "PAGE 2 COLUMN 2 CONTENT 000123" in xml


def test_word_continuation_requires_matching_context(tmp_path):
    from copy import deepcopy
    source, options = generate_sample(tmp_path, "word-basic")
    _, ir = office._extract(source, options, tmp_path, lambda *args: None, lambda: False)
    second = ir.tables[0].model_copy(deep=True)
    second.id = "continuation"
    for cell in second.cells:
        cell.id = "continuation-" + cell.id
    ir.tables.append(second)
    ir.blocks.append(Block(id="continuation-block", kind="table", table_id=second.id, source=second.source))
    merged = office._word_output_tables(ir, {"merge_continuation_tables": True})
    assert len(merged) == 1 and merged[0].row_count == 9
    assert len(office._word_output_tables(ir, {})) == 2
    second.title = "完全不同的单位和业务上下文"
    assert len(office._word_output_tables(ir, {"merge_continuation_tables": True})) == 2


def test_cancel_terminates_only_its_process_tree():
    import subprocess
    import sys
    import time
    own = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(30)"], **office._spawn_kwargs())
    companion = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(30)"], **office._spawn_kwargs())
    started = time.monotonic()
    try:
        with pytest.raises(Cancelled):
            office._wait(own, lambda: time.monotonic()-started > .2, timeout=5)
        assert own.poll() is not None
        assert companion.poll() is None
        assert time.monotonic()-started < 5
    finally:
        office._terminate(own)
        office._terminate(companion)


def test_real_image_block_exports_inside_page_width(tmp_path):
    from PIL import Image
    image = tmp_path / "figure.png"
    Image.new("RGB", (300, 160), "teal").save(image)
    ir = DocumentIR(source_type="pdf", blocks=[Block(id="text", text="Native text remains editable"),
                                              Block(id="figure", kind="image", style={"image_path": str(image), "width_pt": 999})])
    output = office.write_docx(ir, tmp_path / "mixed.docx", {})
    document = Document(output)
    assert len(document.inline_shapes) == 1
    assert document.inline_shapes[0].width <= document.sections[0].page_width-document.sections[0].left_margin-document.sections[0].right_margin
    assert document.paragraphs[0].text == "Native text remains editable"


@pytest.mark.parametrize("case", ["word-basic", "excel-basic"])
def test_encrypted_office_real_password_handling(tmp_path, case):
    from msoffcrypto.format.ooxml import OOXMLFile
    source, options = generate_sample(tmp_path, case)
    encrypted = tmp_path / ("encrypted" + source.suffix)
    with source.open("rb") as original, encrypted.open("wb") as output:
        OOXMLFile(original).encrypt("synthetic-document-password", output)
    with pytest.raises(ToolError) as missing:
        office._normalize(encrypted, {}, tmp_path, lambda: False)
    assert missing.value.code == "PASSWORD_REQUIRED"
    with pytest.raises(ToolError) as wrong:
        office._normalize(encrypted, {"password": "wrong"}, tmp_path, lambda: False)
    assert wrong.value.code == "INVALID_PASSWORD"
    decrypted = office._normalize(encrypted, {"password": "synthetic-document-password"}, tmp_path, lambda: False)
    assert zipfile.is_zipfile(decrypted)
    assert source.read_bytes() == decrypted.read_bytes()


def run_legacy_smoke(folder):
    import json
    from pypdf import PdfReader
    folder.mkdir(parents=True, exist_ok=True)
    results = []
    for case, suffix in [("word-basic", ".doc"), ("excel-basic", ".xls")]:
        native, options = generate_sample(folder / "sources", case)
        legacy = folder / ("legacy" + suffix)
        office.render_office(native, legacy, {}, folder / "make-legacy", lambda: False)
        assert legacy.read_bytes()[:8] == bytes.fromhex("D0CF11E0A1B11AE1")
        for operation in (["word_to_pdf", "word_to_excel"] if suffix == ".doc" else ["excel_to_pdf", "excel_to_word"]):
            result = office.convert_office(legacy, operation, {}, folder / operation, lambda *args: None, lambda: False)
            output = result.files[0]["path"]
            assert output.is_file()
            if output.suffix == ".pdf":
                assert len(PdfReader(output).pages) >= 1
            results.append({"input": str(legacy), "operation": operation, "output": str(output), "issues": [i.code for i in result.ir.issues]})
            print(json.dumps(results[-1], ensure_ascii=False), flush=True)
    (folder / "legacy-report.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    return results


def run_options_smoke(folder):
    import json
    from pypdf import PdfReader
    results = []
    for case, operation, options in [("excel-wide", "excel_to_pdf", {"paper": "A3", "orientation": "landscape", "print_mode": "fit_width"}),
                                     ("excel-hidden", "excel_to_pdf", {"include_hidden": True, "sheets": ["隐藏工作表"]}),
                                     ("word-basic", "word_to_pdf", {}), ("excel-basic", "excel_to_pdf", {})]:
        source, _ = generate_sample(folder / "sources", case)
        original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        ir = None
        if case in ("word-basic", "excel-basic"):
            _, ir = office._extract(source, options, folder, lambda *args: None, lambda: False)
            cell = next(c for c in ir.tables[0].cells if c.display_text == "00001234567890123456789")
            cell.display_text = cell.value = "CORRECTED9988"
            cell.resolution, cell.source.method = "manually_confirmed", "manual"
        result = office.convert_office(source, operation, options, folder / case, lambda *args: None, lambda: False, ir=ir)
        output = result.files[0]["path"]
        pdf = PdfReader(output)
        text = "".join("".join(p.extract_text() or "" for p in pdf.pages).split())
        if ir:
            assert "CORRECTED9988" in text
            assert "00001234567890123456789" not in text
        if case == "excel-wide":
            assert float(pdf.pages[0].mediabox.width) > float(pdf.pages[0].mediabox.height)
            assert len(pdf.pages) == 1
        if case == "excel-hidden":
            assert "不得默认导出" in text
            assert "产品1" not in text
        assert hashlib.sha256(source.read_bytes()).hexdigest() == original_hash
        row = {"case": case, "options": options, "manual_revision": ir is not None, "status": "passed", "pages": len(pdf.pages), "output": str(output)}
        results.append(row)
        print(json.dumps(row, ensure_ascii=False), flush=True)
    (folder / "options-report.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    return results


def reconcile_smoke_reports(original_folder, word_folder, report_path):
    """Recheck existing PDF contents and combine final renderer replacement runs."""
    import json
    rows = json.loads((original_folder / "smoke-report.json").read_text(encoding="utf-8"))
    replacements = {(r["case"], r["operation"]): r for r in json.loads((word_folder / "smoke-report.json").read_text(encoding="utf-8"))}
    for index, row in enumerate(rows):
        replacement = replacements.get((row["case"], row["operation"]))
        if replacement:
            rows[index] = replacement
            continue
        if row["operation"].endswith("_to_pdf") and row.get("output"):
            source = next((original_folder / "sources").glob(row["case"] + " *"))
            options = {"include_headers_footers": True} if row["case"] == "word-header" else {"range": "所选区域"} if row["case"] == "excel-named" else {}
            _, ir = office._extract(source, options, original_folder, lambda *args: None, lambda: False)
            output = Path(row["output"])
            response = list(output.parent.glob("office-response-*.json"))
            if response:
                metadata = json.loads(response[-1].read_text(encoding="utf-8"))
                office._apply_calc_metadata(ir, metadata)
            office._validate_pdf_content(ir, output, options)
            row["issues"] = [{"code": i.code, "message": i.message, "severity": i.severity} for i in ir.issues]
            row["content_check"] = ir.engine_manifest["content_check"]
    summary = {"generated_samples": 20, "conversion_runs": len(rows), "successful_artifacts": sum(r["status"] == "passed" for r in rows),
               "artifact_failures": sum(r["status"] != "passed" for r in rows),
               "note": "产物成功与内容无损不同；问题列表和原文覆盖检查单独记录。自建样本不代表真实业务文件验收。", "results": rows}
    report_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in summary.items() if k != "results"}, ensure_ascii=False), flush=True)
    return summary


def run_real_smoke(folder, operations=None):
    """Opt-in real engine benchmark; JSON distinguishes errors from quality issues."""
    import json
    import time
    from pypdf import PdfReader
    rows = []
    for case in CASES:
        source, options = generate_sample(folder / "sources", case)
        for operation in (["word_to_pdf", "word_to_excel"] if case.startswith("word") else ["excel_to_pdf", "excel_to_word"]):
            if operations and operation not in operations:
                continue
            start = time.monotonic()
            target = folder / case / operation
            try:
                result = office.convert_office(source, operation, options, target, lambda *args: None, lambda: False)
                file = next(f["path"] for f in result.files if f["role"] == "result")
                pages = len(PdfReader(file).pages) if file.suffix == ".pdf" else None
                text = "\n".join(p.extract_text() or "" for p in PdfReader(file).pages) if pages else ""
                key_matches = "00001234567890123456789" in "".join(text.split()) if pages else True
                row = {"case": case, "operation": operation, "status": "passed" if key_matches else "content_mismatch", "output": str(file), "pages": pages,
                       "issues": [{"code": i.code, "message": i.message} for i in result.ir.issues], "seconds": round(time.monotonic()-start, 2)}
            except Exception as exc:
                row = {"case": case, "operation": operation, "status": "failed", "error": str(exc), "seconds": round(time.monotonic()-start, 2)}
            rows.append(row)
            print(json.dumps(row, ensure_ascii=False), flush=True)
            (folder / "smoke-report.json").write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    return rows


if __name__ == "__main__":
    import sys
    run_real_smoke(Path(sys.argv[1]))
