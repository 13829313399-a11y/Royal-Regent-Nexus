"""Loopback HTTP acceptance on pytest-owned SQLite, never a running service.

100/500 dispatched orders and one complete nonempty business chain; 5/10
independent cookie sessions. This bounded read/export workload is not a
production SLA, full-history capacity approval or concurrent financial writing.
"""
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager, ExitStack
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from threading import Thread, Barrier
from time import monotonic, perf_counter, sleep
from types import SimpleNamespace
import json
import socket
from zipfile import ZipFile
from hashlib import sha256

import httpx
import pytest
import uvicorn
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.services.auth import get_current_user
from test_cutting_p5b_acceptance import (
    client, flow, role_flow, prepare_closed_order, sheet_records, BASE,
)
from test_cutting_orders import Line, ledger, DispatchIn
from test_cutting_p5c import evidence, logical_digest
from test_cutting_sqlite_capacity import metrics


@contextmanager
def isolated_http(p, root):
    """No URL/config input: bind a new exclusive loopback socket for this app."""
    database = Path(p.engine.url.database).resolve()
    if not database.is_relative_to(root.resolve()) or database.name != 'cutting.db':
        raise ValueError('HTTP acceptance requires a pytest-owned cutting.db')
    if get_current_user in p.app.dependency_overrides:
        raise ValueError('HTTP acceptance requires real authentication')
    sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    server = None
    thread = None
    try:
        # Do not set SO_REUSEADDR: never share a local business service port.
        sock.bind(('127.0.0.1', 0))
        port = sock.getsockname()[1]
        sock.listen(128)
        server = uvicorn.Server(uvicorn.Config(
            p.app, host='127.0.0.1', port=port, lifespan='off',
            access_log=False, log_level='error', timeout_graceful_shutdown=5,
        ))
        thread = Thread(target=server.run, kwargs={'sockets': [sock]}, daemon=True)
        thread.start()
        deadline = monotonic() + 10
        while not server.started:
            if not thread.is_alive() or monotonic() >= deadline:
                raise RuntimeError('Isolated HTTP server did not start')
            sleep(.02)
        yield f'http://127.0.0.1:{port}'
    finally:
        if server is not None:
            server.should_exit = True
        if thread is not None and thread.ident is not None:
            thread.join(15)
        sock.close()
        if thread is not None and thread.is_alive():
            raise RuntimeError('Isolated HTTP server did not stop')


@pytest.mark.parametrize('outside,override', [(True, False), (False, True)])
def test_http_guard_rejects_unsafe_fixture_before_binding(tmp_path, monkeypatch, outside, override):
    database = tmp_path.parent/'business.db' if outside else tmp_path/'cutting.db'
    p = SimpleNamespace(
        engine=SimpleNamespace(url=SimpleNamespace(database=str(database))),
        app=SimpleNamespace(dependency_overrides={get_current_user: object()} if override else {}),
    )
    def forbidden_socket(*args, **kwargs):
        pytest.fail('Unsafe fixture must be rejected before opening a socket')
    monkeypatch.setattr(socket, 'socket', forbidden_socket)
    with pytest.raises(ValueError):
        with isolated_http(p, tmp_path):
            pytest.fail('Unsafe fixture was accepted')


def add_dispatched_orders(p, count):
    with Session(p.engine) as db:
        for i in range(1, count):
            row = Line(id=f'http-capacity-{i}', factory_id='huakang-c',
                       customer_code='synthetic', customer_name='合成 HTTP 样本',
                       identity_key=f'http-capacity-{i}', reference_no=f'HTTP-{i:06}',
                       product_no='00012', quantity=Decimal('100'),
                       shipped_quantity=0, status='active', version=1, revision=1,
                       data={'quantity': '100', 'requested_ship_date': '2026-10-30'},
                       created_at='2026-10-09', updated_at='2026-10-09')
            db.add(row)
            db.flush()
            ledger.dispatch(db, row, DispatchIn(expected_revision=1, recipients=['cutting']), 'synthetic')
        db.commit()


def login_http(session, role):
    response = session.post('/api/auth/login', json={
        'username': 'p5b-' + role, 'password': 'SyntheticPass123!',
    })
    assert response.status_code == 200
    assert response.json()['identity']['identity_mode'] == 'v2'
    assert session.cookies and 'HttpOnly' in response.headers['set-cookie']


def validate_workbook(response, multiplier=1, expected_delivery=None, *, full_readback=False):
    assert response.headers['cache-control'] == 'no-store'
    collected, covered = [], []
    def inspect_book(raw, volume, coverage):
        book = load_workbook(BytesIO(raw), read_only=True)
        try:
            if full_readback:
                embedded = [dict(zip(('volume','sheet','source','first','last','rows','total'),row))
                            for row in list(book['完整性清单'].values)[2:]]
                assert embedded == coverage
            for part in coverage:
                if part['volume'] != volume:
                    continue
                assert book[part['sheet']].max_row == part['rows']+2
                if full_readback:
                    sheet = book[part['sheet']]
                    sheet.reset_dimensions()
                    assert sum(1 for _ in sheet.iter_rows(min_row=3,values_only=True)) == part['rows']
                covered.append(part)
                if part['source']=='费用汇总':
                    collected.extend(sheet_records(book,part['sheet']))
        finally:
            book.close()
    if response.headers.get('content-type','').split(';')[0] == 'application/zip':
        with ZipFile(BytesIO(response.content)) as archive:
            manifest = json.loads(archive.read('manifest.json'))
            files, coverage = manifest['files'], manifest['coverage']
            assert manifest['token']==response.headers['x-cutting-report-token']
            assert len({f['name'] for f in files})==len(files)
            assert sorted(archive.namelist())==sorted(['manifest.json']+[f['name'] for f in files])
            if expected_delivery is not None:
                assert expected_delivery==dict(format='zip',volumes=len(files),manifest=coverage)
            for volume,item in enumerate(files,1):
                raw = archive.read(item['name'])
                assert len(raw)==item['bytes'] and sha256(raw).hexdigest()==item['sha256']
                inspect_book(raw,volume,coverage)
    else:
        book = load_workbook(BytesIO(response.content),read_only=True)
        try:
            coverage = [dict(zip(('volume','sheet','source','first','last','rows','total'),row))
                        for row in list(book['完整性清单'].values)[2:]]
        finally:
            book.close()
        if expected_delivery is not None:
            assert expected_delivery==dict(format='xlsx',volumes=1,manifest=coverage)
        inspect_book(response.content,1,coverage)
    assert len(covered)==len(coverage)==len({(r['volume'],r['sheet']) for r in covered})
    for source in {part['source'] for part in coverage}:
        parts = sorted((p for p in coverage if p['source']==source),key=lambda p:p['first'])
        cursor,total = 1,parts[0]['total']
        for part in parts:
            assert part['total']==total
            if total:
                assert part['first']==cursor and part['last']-part['first']+1==part['rows']>0
                cursor = part['last']+1
            else:
                assert part['first']==part['last']==part['rows']==0 and len(parts)==1
        assert cursor-1==total
    assert len(collected)==1
    value = collected[0]
    assert tuple(value[k] for k in (
        '币种', '本期正常费用', '期末累计完工费用', '期末在制费用', '本期损耗费用',
    )) == ('HKD', *(f'{amount*multiplier:.2f}' for amount in (Decimal('8'), Decimal('4.8'), Decimal('3.2'), Decimal('2'))))


def projection(name, value):
    """Compare relevant numeric/source facts, excluding generated timestamps."""
    if name == 'planning':
        return value['total'], value['stats']
    if name == 'tasks':
        return value['day'], value['total'], value['stats']
    if name == 'inventory':
        return value['total'], value['units'], [
            (r['issue_id'], r['received_quantity'], r['remaining_quantity']) for r in value['data']
        ]
    if name == 'costs':
        return value['total'], value['currencies'], value['pending_pricing']
    if name == 'reports':
        return value['summary']
    if name == 'period':
        return value['state'], value['event_id'], value['data']['snapshot']['costs']
    if name == 'settlement':
        return value['state'], value['event_id'], value['data']['snapshot']['total']
    if name == 'payments':
        return value['totals']
    raise ValueError('Unknown HTTP workload')


def endpoint_metrics(samples, elapsed, failures):
    result = metrics(samples, elapsed, 0) if samples else {'requests': 0}
    result.pop('process_cpu_seconds', None)
    if samples:
        result['http_failures'] = result.pop('failures')
        result['http_failure_rate'] = result.pop('failure_rate')
    result['validation_failures'] = len(failures)
    result['passed'] = bool(samples) and not failures and result.get('http_failures') == 0
    return result


def test_http_200_with_invalid_content_is_failed_evidence():
    result = endpoint_metrics([{'seconds': .1, 'status': 200, 'bytes': 12}], 1,
                              [{'type': 'AssertionError'}])
    assert result['http_failure_rate'] == 0
    assert result['validation_failures'] == 1 and not result['passed']
    assert not endpoint_metrics([], 1, [])['passed']


@pytest.mark.parametrize('count,workers', [(100, 5), (500, 10)])
def test_cookie_http_read_export_baseline(role_flow, tmp_path, count, workers):
    p = role_flow
    prepare_closed_order(p)
    add_dispatched_orders(p, count)
    endpoints = [
        ('planning', '/workspace/planning', {}),
        ('tasks', '/workspace/tasks', {'day': '2026-10-15'}),
        ('inventory', '/workspace/inventory', {}),
        ('costs', '/workspace/costs', {}),
        ('reports', '/orders/line-1/reports', {'history_page': 1}),
        ('period', '/periods/2026-10', {}),
        ('settlement', '/settlements/2026-10', {'document_id': p.one}),
        ('payments', '/payments/2026-11', {}),
    ]
    samples = []
    failures = []
    durations = []
    with isolated_http(p, tmp_path) as url, ExitStack() as stack:
        # Each worker owns its connection pool and cookie jar. Login/setup is
        # deliberately outside measured business reads and the unchanged check.
        sessions = [stack.enter_context(httpx.Client(
            base_url=url, timeout=30, trust_env=False, follow_redirects=False,
        )) for _ in range(workers)]
        for session in sessions:
            login_http(session, 'reviewer')
        # A live clerk cookie must not gain price/finance visibility from load.
        clerk = stack.enter_context(httpx.Client(base_url=url, timeout=30, trust_env=False))
        login_http(clerk, 'clerk')
        assert clerk.get(BASE+'/workspace/costs', params={'factory_id': 'huakang-c'}).status_code == 403
        assert sessions[0].get(BASE+'/workspace/planning', params={'factory_id': 'huakang-d'}).status_code == 422
        expected = {}
        for name, path, params in endpoints:
            response = sessions[0].get(BASE+path, params={'factory_id': 'huakang-c', **params})
            assert response.status_code == 200, (name, response.status_code)
            expected[name] = projection(name, response.json())
        assert expected['planning'][0] == count
        assert expected['tasks'][1] == 2
        assert expected['inventory'][0] == 2
        assert expected['reports']['completed'] == 25
        assert expected['period'][0] == 'locked'
        body = {'factory_id': 'huakang-c', 'kind': 'period', 'month': '2026-10',
                'event_id': expected['period'][1]}
        before = logical_digest(p.engine.url.database)
        barrier = Barrier(workers, timeout=30)

        def worker(index):
            session = sessions[index]
            local, errors = [], []
            barrier.wait()
            start = perf_counter()
            deadline = start + 30
            cycles = 0
            while cycles == 0 or perf_counter() < deadline:
                for name, path, params in endpoints:
                    began = perf_counter()
                    status, size = 0, 0
                    try:
                        response = session.get(BASE+path, params={'factory_id': 'huakang-c', **params})
                        status, size = response.status_code, len(response.content)
                        seconds = perf_counter()-began
                        assert status == 200
                        assert projection(name, response.json()) == expected[name]
                    except Exception as error:
                        seconds = perf_counter()-began
                        errors.append({'worker': index, 'endpoint': name, 'type': type(error).__name__})
                    local.append({'endpoint': name, 'seconds': seconds, 'status': status, 'bytes': size})
                token = None
                for name in ('export_preview', 'export_download'):
                    if name == 'export_download' and token is None:
                        break
                    began = perf_counter()
                    status, size = 0, 0
                    try:
                        response = session.post(BASE+'/exports/'+('preview' if name == 'export_preview' else 'download'),
                                                json=body if token is None else {**body, 'token': token})
                        status, size = response.status_code, len(response.content)
                        seconds = perf_counter()-began
                        assert status == 200
                        if name == 'export_preview':
                            token = response.json()['token']
                        else:
                            validate_workbook(response)
                    except Exception as error:
                        seconds = perf_counter()-began
                        errors.append({'worker': index, 'endpoint': name, 'type': type(error).__name__})
                    local.append({'endpoint': name, 'seconds': seconds, 'status': status, 'bytes': size})
                cycles += 1
                if errors:
                    break
            return local, errors, perf_counter()-start, cycles

        start = perf_counter()
        with ThreadPoolExecutor(max_workers=workers) as pool:
            results = list(pool.map(worker, range(workers)))
        elapsed = perf_counter()-start
        for rows, errors, duration, cycles in results:
            samples.extend(rows)
            failures.extend(errors)
            durations.append({'seconds': duration, 'cycles': cycles})
        unchanged = logical_digest(p.engine.url.database) == before
    grouped = {}
    for name in [n for n, _, _ in endpoints] + ['export_preview', 'export_download']:
        values = [s for s in samples if s['endpoint'] == name]
        grouped[name] = endpoint_metrics(values, elapsed, [f for f in failures if f['endpoint'] == name])
        # No process CPU estimate: server and HTTP clients share this process.
    report = dict(
        orders=count, workers=workers, minimum_window_seconds=30, elapsed_seconds=elapsed,
        worker_runs=durations, http_requests=len(samples), validation_failures=failures,
        per_endpoint=grouped, all_database_rows_unchanged=unchanged,
        fixture=dict(dispatched=count, full_business_chains=1, published_tasks=2, signed_issues=2,
                     authentication='real V2 cookie sessions, same synthetic reviewer, separate sessions'),
        scope='loopback TCP HTTP; bounded read/export mix; excludes sustained writes, proxies, attachments and target hardware',
        production_capacity_accepted=False, release_ready=False,
    )
    (evidence(tmp_path)/f'http-capacity-{count}-{workers}.json').write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    assert unchanged
    assert not failures, failures
    assert all(v['requests'] >= workers for v in grouped.values())
    assert all(r['seconds'] >= 30 for r in durations)
