import importlib
import json
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as WorksheetImage
from PIL import Image

from app.services.internal_quote_calculator import CalculationInputError, calculate_section
from app.services.internal_quote_prefill import normalize_engineering_mold_identities, prefill_molding_from_engineering
from test_internal_quote_api import create_payload, login, make_client

SNAPSHOT = {"material_prices": {"ABS|750SW": "8.5", "1#PP|7032 E3": "6.8", "透明ABS|clear": "12"}, "legacy_material_prices": {"ABS": "1"}, "machine_prices": [{"range": "5A", "shift_price_hkd": "1000"}]}
ROW = {"mold_no": "M1", "item": "外壳", "material": "ABS", "grade": "750SW", "machine_code": "5A", "net_weight_g": 100, "sets": 1, "target_output": 1000, "quantity": 1}


def test_missing_draft_fields_block_but_invalid_numeric_fields_reject_even_without_prices():
    result = calculate_section("molding", {"injection_lines": [{**ROW, "grade": "", "sets": 0, "target_output": "", "net_weight_g": 0}]}, SNAPSHOT, "ref")
    assert result["status"] == "blocked"
    assert {error["field"] for error in result["validation_errors"]} >= {"grade", "sets", "target_output", "net_weight_g"}
    for value in (-1, "not a number", "NaN"):
        with pytest.raises(CalculationInputError):
            calculate_section("molding", {"injection_lines": [{**ROW, "grade": "unknown", "target_output": value}]}, SNAPSHOT, "ref")


def test_material_equivalence_is_explicit_and_unknown_grade_never_falls_back():
    result = calculate_section("molding", {"injection_lines": [{**ROW, "material": "PP", "grade": "7032 E3"}]}, SNAPSHOT, "ref")
    assert result["status"] == "valid"
    assert result["line_breakdown"][0]["material_price_hkd_lb"] == "6.8000"
    for material, grade in (("ABS", "clear"), ("ABS", "unknown"), ("透明ABS", "750SW")):
        result = calculate_section("molding", {"injection_lines": [{**ROW, "material": material, "grade": grade}]}, SNAPSHOT, "ref")
        assert result["status"] == "blocked"
        assert result["validation_errors"][0]["field"] == "grade"


def test_stable_ids_preserve_children_and_cost_once_after_rename():
    source = normalize_engineering_mold_identities({"molds": [{"mold_no": "M1", "item": "外壳", "net_weight_g": 100, "parts": [{"name": "左", "unit_net_weight_g": 40}, {"name": "右", "unit_net_weight_g": 60}]}]})
    linked = prefill_molding_from_engineering(source, {"injection_lines": [deepcopy(ROW)]})
    before = calculate_section("molding", linked, SNAPSHOT, "ref")
    source["molds"][0]["mold_no"] = "renamed"
    source["molds"][0]["parts"].reverse()
    relinked = prefill_molding_from_engineering(source, linked)
    assert len(relinked["injection_lines"]) == 1
    assert relinked["injection_lines"][0]["line_id"] == linked["injection_lines"][0]["line_id"]
    assert relinked["injection_lines"][0]["engineering_parts"] == source["molds"][0]["parts"]
    assert calculate_section("molding", relinked, SNAPSHOT, "ref")["totals"] == before["totals"]
    assert calculate_section("molding", prefill_molding_from_engineering({"molds": []}, relinked), SNAPSHOT, "ref")["status"] == "blocked"


def test_detached_suppression_survives_save_reload_rename_delete_and_clone(monkeypatch):
    from test_internal_quote_alternatives import fork

    with make_client(monkeypatch) as client:
        login(client, "detach-owner", "sales_customer_owner", "sales-business")
        created = client.post("/api/internal-quotes", json={**create_payload(suffix="DETACH", participating_sections=["engineering", "sales", "assembly", "molding"]), "workflow_mode": "direct_output"})
        assert created.status_code == 201, created.text
        base = f"/api/internal-quotes/{created.json()['id']}"

        def read_sections():
            return {row["department"]: row for row in client.get(base).json()["sections"]}

        def save(payloads):
            current = read_sections()
            response = client.put(base + "/sections/save-all", json={"sections": [
                {"section_code": code, "revision": current[code]["revision"], "payload": payload}
                for code, payload in payloads.items()
            ]})
            assert response.status_code == 200, response.text
            return read_sections()

        engineering = {"molds": [{"mold_id": "m1", "mold_no": "M1", "item": "外壳", "quantity": 1, "net_weight_g": 100, "cost_rmb": 10, "parts": [{"part_id": "p1", "name": "左"}]}]}
        sections = save({"engineering": engineering, "molding": {"injection_lines": [deepcopy(ROW)]}})
        detached = sections["molding"]["payload"]["injection_lines"][0]
        detached.update(engineering_sync_disabled=True, engineering_synced_fields=[], engineering_parts=[], item="独立外壳", net_weight_g=77)
        detached.pop("cost_scope", None)
        expected = deepcopy(detached)
        save({"molding": {"injection_lines": [detached]}})
        engineering["molds"][0].update(item="改名", mold_no="M2", net_weight_g=999, parts=[{"part_id": "p1", "name": "新部件"}])
        sections = save({"engineering": engineering, "molding": {"injection_lines": [detached]}})
        assert sections["molding"]["payload"]["injection_lines"] == [expected]
        assert sections["molding"]["calculation_status"] == "valid"
        copied = fork(client, client.get(base).json())
        copied_molding = next(row for row in copied["sections"] if row["department"] == "molding")
        assert copied_molding["payload"]["injection_lines"] == [expected]
        assert copied_molding["calculation_status"] == "valid"
        sections = save({"engineering": {"molds": []}, "molding": {"injection_lines": [detached]}})
        assert sections["molding"]["payload"]["injection_lines"] == [expected]
        assert sections["molding"]["calculation_status"] == "valid"
        sections = save({"engineering": {"molds": [{"mold_id": "m-new", "mold_no": "NEW", "item": "新模具", "quantity": 1, "cost_rmb": 10}]}})
        rows = sections["molding"]["payload"]["injection_lines"]
        assert len(rows) == 2
        assert rows[0]["line_id"] == "injection:m-new"
        assert rows[1] == expected


@pytest.mark.parametrize("manual_same_source", [False, True])
def test_genuine_legacy_detachment_whole_save_adopts_identity_without_replacement(monkeypatch, manual_same_source):
    from test_internal_quote_history import put_source

    with make_client(monkeypatch) as client:
        login(client, "detach-legacy", "sales_customer_owner", "sales-business")
        created = client.post("/api/internal-quotes", json={**create_payload(suffix="DETACH-LEGACY", participating_sections=["engineering", "sales", "assembly", "molding"]), "workflow_mode": "direct_output"})
        assert created.status_code == 201, created.text
        base = f"/api/internal-quotes/{created.json()['id']}"
        put_source(created.json()["id"], {
            "engineering": {"molds": [{"mold_no": "M1", "item": "外壳", "quantity": 1, "net_weight_g": 100, "cost_rmb": 10}]},
            "molding": {"injection_lines": [{**ROW, "engineering_source_key": "mold-no:M1#1"}]},
        })
        sections = {row["department"]: row for row in client.get(base).json()["sections"]}
        engineering = sections["engineering"]["payload"]
        detached = sections["molding"]["payload"]["injection_lines"][0]
        assert "mold_id" not in engineering["molds"][0]
        assert "engineering_mold_id" not in detached
        assert "line_id" not in detached
        detached.update(engineering_sync_disabled=True, engineering_synced_fields=[], engineering_parts=[], item="独立", net_weight_g=77)
        manual = {**ROW, "line_id": "manual", "net_weight_g": 31, "engineering_sync_disabled": False, "remark": "手工补充"}
        submitted_rows = [detached, manual] if manual_same_source else [detached]

        def save_both(engineering_payload, molding_payload):
            current = {row["department"]: row for row in client.get(base).json()["sections"]}
            response = client.put(base + "/sections/save-all", json={"sections": [
                {"section_code": code, "revision": current[code]["revision"], "payload": payload}
                for code, payload in (("engineering", engineering_payload), ("molding", molding_payload))
            ]})
            assert response.status_code == 200, response.text
            return {row["department"]: row for row in client.get(base).json()["sections"]}

        sections = save_both(engineering, {"injection_lines": submitted_rows})
        adopted = sections["engineering"]["payload"]
        mold = adopted["molds"][0]
        assert mold["legacy_source_key"] == "mold-no:M1#1"
        rows = sections["molding"]["payload"]["injection_lines"]
        assert len(rows) == len(submitted_rows)
        assert rows[0] == {**detached, "engineering_source_key": f"mold-id:{mold['mold_id']}", "engineering_mold_id": mold["mold_id"], "line_id": f"injection:{mold['mold_id']}"}
        if manual_same_source:
            assert rows[1] == manual
        assert sections["molding"]["calculation_status"] == "valid"
        saved_molding = deepcopy(sections["molding"]["payload"])
        mold.update(item="改名", mold_no="renamed", net_weight_g=999)
        sections = save_both(adopted, saved_molding)
        assert sections["molding"]["payload"] == saved_molding
        sections = save_both({"molds": []}, saved_molding)
        assert sections["molding"]["payload"] == saved_molding
        assert sections["molding"]["calculation_status"] == "valid"


@pytest.mark.parametrize("ambiguity", ["duplicate_legacy_rows", "duplicate_source_aliases", "stable_and_legacy_rows"])
def test_ambiguous_legacy_association_blocks_instead_of_generating_prices(ambiguity, monkeypatch):
    source = normalize_engineering_mold_identities({"molds": [{"mold_no": "M1", "item": "外壳"}]})
    detached = {**ROW, "line_id": "old-1", "engineering_source_key": "mold-no:M1#1", "engineering_sync_disabled": True}
    rows = [detached]
    if ambiguity == "duplicate_source_aliases":
        source["molds"].append({**source["molds"][0], "mold_id": "different-id"})
    elif ambiguity == "stable_and_legacy_rows":
        mold_id = source["molds"][0]["mold_id"]
        rows.append({**ROW, "line_id": "replacement", "engineering_mold_id": mold_id, "engineering_source_key": f"mold-id:{mold_id}"})
    else:
        rows.append({**detached, "line_id": "old-2"})
    result = prefill_molding_from_engineering(source, {"injection_lines": rows})
    assert len(result["injection_lines"]) == len(rows)
    assert all(row["engineering_linkage_conflict"] for row in result["injection_lines"])
    calculated = calculate_section("molding", result, SNAPSHOT, "ref")
    assert calculated["status"] == "blocked"
    assert calculated["validation_errors"][0]["field"] == "engineering_source_key"
    assert "旧模具关联" in calculated["validation_errors"][0]["message"]
    assert all(line["amount_hkd"] is None for line in calculated["line_breakdown"])
    # Existing mutable records with old valid totals must also be refused at
    # release without rewriting their payloads merely to detect the conflict.
    from fastapi import HTTPException
    service = importlib.import_module("app.services.internal_quote")
    monkeypatch.setattr(service, "_find_reference_set", lambda *_: SimpleNamespace(id="ref", snapshot_json=json.dumps(SNAPSHOT)))
    stored = [SimpleNamespace(department="engineering", is_required=True, payload_json=json.dumps(source)),
              SimpleNamespace(department="molding", is_required=True, status="draft", payload_json=json.dumps({"injection_lines": rows}), calculation_json=json.dumps({"totals": {"total_hkd": "1"}}))]
    before = [section.payload_json for section in stored]
    with pytest.raises(HTTPException) as error:
        service.ensure_current_molding_pricing(None, SimpleNamespace(module_version="v4", final_release_status="", factory_id="huaxing"), stored)
    assert error.value.status_code == 409
    assert error.value.detail["validation_errors"][0]["field"] == "engineering_source_key"
    assert [section.payload_json for section in stored] == before
    source["molds"] = source["molds"][:1]
    result["injection_lines"] = result["injection_lines"][:1]
    resolved = prefill_molding_from_engineering(source, result)
    assert "engineering_linkage_conflict" not in resolved["injection_lines"][0]
    assert calculate_section("molding", resolved, SNAPSHOT, "ref")["status"] == "valid"
    manual = {**ROW, "engineering_sync_disabled": True}
    assert len(prefill_molding_from_engineering(source, {"injection_lines": [manual]})["injection_lines"]) == 2


def test_detached_component_history_rebinds_suppression_identity(monkeypatch):
    from test_internal_quote_history import catalog, create, parts, put_source, source_ref

    with make_client(monkeypatch) as client:
        login(client, "detach-history", "sales_customer_owner", "sales-business", "huakang-b")
        source = create(client, "DETACHED-SOURCE", jp=True)
        put_source(source["id"], {
            "engineering": {"molds": [{"mold_id": "m1", "mold_no": "M1", "item": "外壳", "quantity": 1, "cost_rmb": 10, "pricing_component_id": "component-02", "parts": [{"part_id": "p1", "name": "左"}]}]},
            "molding": {"injection_lines": [{**ROW, "line_id": "injection:m1", "engineering_mold_id": "m1", "engineering_source_key": "mold-id:m1", "engineering_sync_disabled": True, "engineering_synced_fields": [], "engineering_parts": [], "net_weight_g": 77, "item": "独立", "pricing_component_id": "component-02"}]},
        })
        source_row = next(row for row in catalog(client) if row["quote_id"] == source["id"])
        copied = create(client, "DETACHED-COMPONENT", jp=True, products=[{"product_name": "新产品", "qty": 100, "pricing_components": ["独立件"], "component_sources": [source_ref(source_row, "component-02")]}])
        mold = parts(copied, "engineering")["molds"][0]
        rows = parts(copied, "molding")["injection_lines"]
        assert len(rows) == 1
        assert mold["mold_id"] != "m1"
        assert rows[0]["engineering_mold_id"] == mold["mold_id"]
        assert rows[0]["engineering_source_key"] == f"mold-id:{mold['mold_id']}"
        assert rows[0]["line_id"] != "injection:m1"
        assert rows[0]["engineering_parts"] == []
        assert rows[0]["engineering_sync_disabled"] is True
        assert rows[0]["item"] == "独立" and rows[0]["net_weight_g"] == 77


def painting_workbook():
    wb = Workbook()
    sheet = wb.active
    sheet.append(["图片", "名称", "位置", "UV", "UV单价", "总报价"])
    sheet.append([None, "外壳", "正面", 2, .52, 1.04])
    sheet.row_dimensions[2].height = 40
    png = BytesIO()
    Image.new("RGB", (20, 20), "red").save(png, "PNG")
    sheet.add_image(WorksheetImage(BytesIO(png.getvalue())), "A2")
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def test_painting_image_import_confirm_preview_export_and_source_delete_keep_manual_reference(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "painting_images", "sales_customer_owner", "sales-business")
        created = client.post("/api/internal-quotes", json={**create_payload(suffix="IMAGE-LINKAGE", participating_sections=["engineering", "sales", "assembly", "painting"]), "workflow_mode": "direct_output"})
        assert created.status_code == 201, created.text
        base = f"/api/internal-quotes/{created.json()['id']}"
        response = client.post(base + "/imports/painting/preview", files={"file": ("painting.xlsx", painting_workbook())})
        assert response.status_code == 201, response.text
        preview = response.json()
        assert len(preview["embedded_images"]) == 1
        confirmed = client.post(base + f"/imports/{preview['batch_id']}/confirm", json={"revision": preview["target_revision"], "mode": "replace"})
        assert confirmed.status_code == 200, confirmed.text
        section = confirmed.json()["section"]
        row = section["payload"]["rows"][0]
        assert row["row_id"]
        image_id = row["image_attachment_ids"][0]
        assert client.get(base + f"/attachments/{image_id}/preview").status_code == 200
        # Export uses only scoped attachments, and emits the actual image bytes.
        db_module = importlib.import_module("app.db")
        models = importlib.import_module("app.models.internal_quote")
        excel = importlib.import_module("app.services.internal_quote_excel")
        with db_module.SessionLocal() as db:
            stored = db.query(models.InternalQuoteSection).filter_by(quote_id=created.json()["id"], department="painting").one()
            attachment = db.get(models.InternalQuoteAttachment, image_id)
            output = Workbook()
            excel._build_painting_image_sheet(output, stored, [attachment])
            stream = BytesIO()
            output.save(stream)
            assert len(load_workbook(BytesIO(stream.getvalue()))["喷油图片明细"]._images) == 1
        manual = {**row, "row_id": "manual-image", "name": "手工引用"}
        manual.pop("import_batch_id")
        section["payload"]["rows"].append(manual)
        saved = client.put(base + "/sections/painting", json={"revision": section["revision"], "payload": section["payload"]})
        assert saved.status_code == 200, saved.text
        attachments = client.get(base + "/attachments?department=painting").json()
        source = next(item for item in attachments if item["is_import_source"])
        deleted = client.delete(base + f"/attachments/{source['id']}", params={"revision": saved.json()["revision"]})
        assert deleted.status_code == 204, deleted.text
        remaining = next(item for item in client.get(base).json()["sections"] if item["department"] == "painting")
        assert len(remaining["payload"]["rows"]) == 1
        assert remaining["payload"]["rows"][0]["row_id"] == "manual-image"
        assert client.get(base + f"/attachments/{image_id}/preview").status_code == 200
        from test_internal_quote_alternatives import fill, fork
        ready = fill(client, client.get(base).json())
        issued = client.post(base + "/direct-issue", json={"revision": ready["header_revision"]})
        assert issued.status_code == 200, issued.text
        download = base + f"/exports/{issued.json()['id']}/download"
        frozen_bytes = client.get(download).content
        assert len(load_workbook(BytesIO(frozen_bytes))["喷油图片明细"]._images) == 1
        frozen = client.get(base).json()
        copied = fork(client, frozen, "version")
        copied_row = next(item for item in copied["sections"] if item["department"] == "painting")["payload"]["rows"][0]
        copied_id = copied_row["image_attachment_ids"][0]
        assert copied_id != image_id
        assert client.get(f"/api/internal-quotes/{copied['id']}/attachments/{copied_id}/preview").status_code == 200
        copied_section = next(item for item in copied["sections"] if item["department"] == "painting")
        copied_section["payload"]["rows"][0]["image_attachment_ids"] = [image_id]
        cross_quote = client.put(f"/api/internal-quotes/{copied['id']}/sections/painting", json={"revision": copied_section["revision"], "payload": copied_section["payload"]})
        assert cross_quote.status_code == 400
        blocked = client.put(base + "/sections/painting", json={"revision": remaining["revision"], "payload": remaining["payload"]})
        assert blocked.status_code == 409
        assert client.get(download).content == frozen_bytes


def test_save_all_merges_drafts_before_linking_and_preserves_atomicity(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "linkage_save", "sales_customer_owner", "sales-business")
        created = client.post("/api/internal-quotes", json=create_payload(suffix="LINKAGE-SAVE", participating_sections=["engineering", "sales", "assembly", "molding"]))
        assert created.status_code == 201, created.text
        base = f"/api/internal-quotes/{created.json()['id']}"
        engineering = {"molds": [{"mold_id": "m1", "mold_no": "M1", "item": "外壳", "material": "ABS", "quantity": 1, "cost_rmb": 10}]}
        drafts = [{"section_code": "engineering", "revision": 1, "payload": engineering}, {"section_code": "molding", "revision": 1, "payload": {"injection_lines": [{**ROW, "grade": "", "target_output": 0}]}}]
        saved = client.put(base + "/sections/save-all", json={"sections": drafts})
        assert saved.status_code == 200, saved.text
        sections = {row["department"]: row for row in saved.json()["sections"]}
        assert sections["molding"]["calculation_status"] == "blocked"
        assert len(sections["molding"]["payload"]["injection_lines"]) == 1
        assert sections["molding"]["calculation"]["validation_errors"]
        engineering["molds"][0]["target_output"] = 1000
        drafts[0].update(revision=sections["engineering"]["revision"])
        drafts[1].update(revision=sections["molding"]["revision"], payload={"injection_lines": [{**sections["molding"]["payload"]["injection_lines"][0], "grade": "750SW"}]})
        saved = client.put(base + "/sections/save-all", json={"sections": drafts})
        assert saved.status_code == 200, saved.text
        sections = {row["department"]: row for row in saved.json()["sections"]}
        assert sections["molding"]["calculation_status"] == "valid"
        assert sections["molding"]["dependency_status"] == "current"
        drafts[0].update(revision=sections["engineering"]["revision"], payload={**engineering, "materials": [{"item": "changed", "quantity": 1, "unit_price_rmb": 1}]})
        drafts[1].update(revision=sections["molding"]["revision"], payload={"injection_lines": [{**ROW, "quantity": -1}]})
        failed = client.put(base + "/sections/save-all", json={"sections": drafts})
        assert failed.status_code == 400
        after = {row["department"]: row for row in client.get(base).json()["sections"]}
        assert after["engineering"]["revision"] == sections["engineering"]["revision"]
        assert after["engineering"]["payload"] == sections["engineering"]["payload"]
        invalid_engineering = deepcopy(sections["engineering"]["payload"])
        invalid_engineering["molds"][0]["net_weight_g"] = -1
        drafts[0].update(payload=invalid_engineering)
        drafts[1].update(payload=sections["molding"]["payload"])
        failed = client.put(base + "/sections/save-all", json={"sections": drafts})
        assert failed.status_code == 400, failed.text
        after = {row["department"]: row for row in client.get(base).json()["sections"]}
        assert after["engineering"]["payload"] == sections["engineering"]["payload"]
        assert after["molding"]["revision"] == sections["molding"]["revision"]


def test_hape_complete_and_stale_summary_never_exposes_partial_as_quote(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "hape_complete", "sales_customer_owner", "sales-business")
        created = client.post("/api/internal-quotes", json=create_payload(suffix="HAPE-COMPLETE", participating_sections=["engineering", "sales", "assembly", "molding"]))
        assert created.status_code == 201, created.text
        quote_id = created.json()["id"]
        base = f"/api/internal-quotes/{quote_id}"
        db_module = importlib.import_module("app.db")
        models = importlib.import_module("app.models.internal_quote")
        service = importlib.import_module("app.services.internal_quote")
        with db_module.SessionLocal() as db:
            quote = db.get(models.InternalQuote, quote_id)
            for section in db.query(models.InternalQuoteSection).filter_by(quote_id=quote_id, is_required=True):
                payload = {"shipping": {"markup_x": "1.12", "misc_ratio": ".02"}} if section.department == "sales" else {"confirmed": True}
                totals = {"total_hkd": "27.248"} if section.department == "molding" else {"hardware_hkd": "18.032", "total_hkd": "18.032"} if section.department == "engineering" else {}
                section.payload_json = json.dumps(payload)
                section.calculation_json = json.dumps({"status": "valid", "totals": totals, "warnings": [], "input_hash": service.content_hash(payload)})
                section.calculation_status = "valid"
                section.dependency_status = "current"
                section.calculation_formula_version = quote.formula_version
                section.calculation_reference_snapshot_id = quote.reference_snapshot_id
                section.filled_at = "2026-10-14"
            db.flush()
            for section in db.query(models.InternalQuoteSection).filter_by(quote_id=quote_id, is_required=True):
                section.dependency_hash = service.content_hash(service._calculation_dependencies(db, quote, section.department))
            db.commit()
        complete = client.get(base + "/summary").json()
        assert complete["pricing_complete"] is True
        assert complete["factory_price_hkd"] == "45.2800"
        assert complete["authoritative_factory_quote_hkd"] == "51.7486"
        blocked_preview = client.post(base + "/cost-preview", json={"drafts": [{"section_code": "molding", "revision": 1, "payload": {"injection_lines": [{**ROW, "grade": "unknown"}]}}]})
        assert blocked_preview.status_code == 200, blocked_preview.text
        assert blocked_preview.json()["preview_factory_price_hkd"] == "18.0320"
        assert blocked_preview.json()["authoritative_factory_quote_hkd"] is None
        assert client.get(base + "/summary").json()["factory_price_hkd"] == "45.2800"
        with db_module.SessionLocal() as db:
            section = db.query(models.InternalQuoteSection).filter_by(quote_id=quote_id, department="molding").one()
            section.calculation_status = "stale"
            db.commit()
        partial = client.get(base + "/summary").json()
        assert partial["pricing_complete"] is False
        assert partial["factory_price_hkd"] == "18.0320"
        assert partial["authoritative_factory_quote_hkd"] is None
        assert partial["incomplete_sections"][0]["section_code"] == "molding"
        preview = client.post(base + "/cost-preview", json={"drafts": [{"section_code": "molding", "revision": 1, "payload": {"injection_lines": [{**ROW, "grade": "unknown"}]}}]})
        assert preview.status_code == 200, preview.text
        assert preview.json()["preview_factory_price_hkd"] == "18.0320"
        assert preview.json()["authoritative_factory_quote_hkd"] is None


@pytest.mark.parametrize("workflow_mode", ["direct_output", "whole_quote_review", "section_review"])
def test_old_valid_mutable_molding_is_revalidated_without_rewriting_at_issue_or_submit(monkeypatch, workflow_mode):
    with make_client(monkeypatch) as client:
        login(client, "old_valid_molding", "sales_customer_owner", "sales-business")
        created = client.post("/api/internal-quotes", json={**create_payload(suffix="OLD-VALID", participating_sections=["engineering", "sales", "assembly", "molding"]), "workflow_mode": workflow_mode})
        assert created.status_code == 201, created.text
        quote_id = created.json()["id"]
        base = f"/api/internal-quotes/{quote_id}"
        from test_internal_quote_alternatives import fill
        fill(client, created.json())
        saved = client.put(base + "/sections/molding", json={"revision": 1, "payload": {"injection_lines": [ROW]}})
        assert saved.status_code == 200, saved.text
        assert saved.json()["calculation_status"] == "valid"
        db_module = importlib.import_module("app.db")
        models = importlib.import_module("app.models.internal_quote")
        service = importlib.import_module("app.services.internal_quote")
        with db_module.SessionLocal() as db:
            quote = db.get(models.InternalQuote, quote_id)
            sections = db.query(models.InternalQuoteSection).filter_by(quote_id=quote_id, is_required=True).all()
            molding = next(row for row in sections if row.department == "molding")
            # Simulate a previously valid v7 record accepted through the former
            # material-only fallback, without changing its stored calculation.
            molding.payload_json = json.dumps({"injection_lines": [{**ROW, "grade": "unknown-old-grade"}]})
            if workflow_mode == "section_review":
                for section in sections:
                    section.status = "approved"
            service._derive_quote_status(db, quote)
            db.commit()
            before = (quote.header_revision, molding.revision, molding.payload_json, molding.calculation_json)
        assert client.get(base).status_code == 200  # Historical record stays readable.
        endpoint = "/direct-issue" if workflow_mode == "direct_output" else "/final-submit"
        rejected = client.post(base + endpoint, json={"revision": before[0]})
        assert rejected.status_code == 409, rejected.text
        assert rejected.json()["detail"]["section_code"] == "molding"
        assert rejected.json()["detail"]["validation_errors"][0]["field"] == "grade"
        with db_module.SessionLocal() as db:
            quote = db.get(models.InternalQuote, quote_id)
            molding = db.query(models.InternalQuoteSection).filter_by(quote_id=quote_id, department="molding").one()
            assert (quote.header_revision, molding.revision, molding.payload_json, molding.calculation_json) == before
