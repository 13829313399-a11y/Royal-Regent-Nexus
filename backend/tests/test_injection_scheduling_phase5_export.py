import hashlib
import hmac
import importlib
import json
import sys
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from test_injection_scheduling_phase4_import import (
    ADMIN_TEST_PASSWORD,
    build_workbook,
    ensure_user,
    login,
    make_client,
)
from test_injection_scheduling_profile_canonical import _huakang_b_fixture

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))


def _segmented_manifest(workbook) -> dict:
    sheet = workbook["_SYSTEM_META"]
    manifest = json.loads(sheet["A2"].value)
    manifest["rows"] = [
        json.loads(sheet.cell(8 + offset, 1).value)
        for offset in range(manifest["task_count"])
    ]
    return manifest


def _take_over_plan(
    client,
    *,
    factory_id: str,
    source: bytes,
    prefix: str,
    business_date: str,
) -> tuple[dict, dict]:
    preview = client.post(
        "/api/injection-scheduling/imports/preview",
        data={"factory_id": factory_id, "expected_revision": "0"},
        files={
            "file": (
                f"{prefix}.xlsx",
                source,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers={"x-request-id": f"{prefix}-preview-0001"},
    )
    assert preview.status_code == 201, preview.text
    batch = preview.json()
    if batch["master_differences"]:
        approval = client.post(
            f"/api/injection-scheduling/imports/{batch['id']}/master-differences/approve",
            json={
                "factory_id": factory_id,
                "expected_revision": batch["revision"],
                "request_id": f"{prefix}-master-approval-0001",
                "reason": "阶段 5 导出往返契约测试主数据复核",
                "differences": [
                    f"{item['entity_type']}:{item['business_key']}"
                    for item in batch["master_differences"]
                ],
            },
        )
        assert approval.status_code == 200, approval.text
        batch = approval.json()
    assert batch["batch_state"] == "PREVIEW_READY", batch
    confirm = client.post(
        f"/api/injection-scheduling/imports/{batch['id']}/confirm",
        json={
            "factory_id": factory_id,
            "expected_revision": batch["revision"],
            "expected_plan_revision": 0,
            "request_id": f"{prefix}-confirm-0001",
            "confirm_mode": "create_draft",
            "business_date": business_date,
            "acknowledged_blocking_issue_ids": [],
            "expected_action_fingerprint": batch["action_fingerprint"],
        },
    )
    assert confirm.status_code == 200, confirm.text
    current = client.get(
        "/api/injection-scheduling/plans/current",
        params={"factory_id": factory_id},
    )
    assert current.status_code == 200, current.text
    return batch, current.json()["plan"]


@pytest.mark.parametrize(
    ("factory_id", "source", "profile_id", "sheet_name", "business_date"),
    [
        (
            "huaxing",
            build_workbook(),
            "isprofile-huaxing-daily-v1",
            "计划表",
            "2026-08-01",
        ),
        (
            "huakang-b",
            _huakang_b_fixture(),
            "isprofile-huakang-b-daily-v1",
            "排期表",
            "2026-08-05",
        ),
    ],
    ids=("huaxing", "huakang-b"),
)
def test_phase5_source_compatible_export_is_signed_audited_and_round_trips(
    monkeypatch,
    factory_id,
    source,
    profile_id,
    sheet_name,
    business_date,
):
    monkeypatch.setenv(
        "INJECTION_SCHEDULING_EXPORT_SIGNING_KEY",
        "phase5-test-signing-key-with-at-least-32-bytes",
    )
    monkeypatch.setenv("INJECTION_SCHEDULING_EXPORT_SIGNING_KEY_ID", "phase5-test-v1")
    prefix = f"phase5-{factory_id}"
    with make_client(monkeypatch) as client:
        login(client, "admin", ADMIN_TEST_PASSWORD)
        _batch, plan = _take_over_plan(
            client,
            factory_id=factory_id,
            source=source,
            prefix=prefix,
            business_date=business_date,
        )
        assert plan["export_profile_id"] == profile_id
        assert plan["export_profile_revision"] == 1
        assert plan["export_binding_source"] == "IMPORT_PROFILE"
        assert plan["calculation_version"]

        export_payload = {
            "factory_id": factory_id,
            "expected_plan_revision": plan["revision"],
            "request_id": f"{prefix}-source-export-0001",
            "export_mode": "SOURCE_COMPATIBLE",
            "profile_id": profile_id,
            "profile_revision": 1,
        }
        exported = client.post(
            f"/api/injection-scheduling/plans/{plan['id']}/exports",
            json=export_payload,
        )
        assert exported.status_code == 200, exported.text
        assert exported.headers["x-export-mode"] == "SOURCE_COMPATIBLE"
        assert exported.headers["x-export-profile-id"] == profile_id
        assert exported.headers["x-export-signing-key-id"] == "phase5-test-v1"
        assert exported.headers["x-idempotent-replay"] == "false"
        assert exported.headers["x-export-sha256"] == hashlib.sha256(
            exported.content
        ).hexdigest()

        workbook = load_workbook(BytesIO(exported.content), data_only=False)
        assert workbook.sheetnames[0] == sheet_name
        assert "_SYSTEM_META" in workbook.sheetnames
        assert workbook["_SYSTEM_META"].sheet_state == "veryHidden"
        assert workbook["_SYSTEM_META"]["A6"].value == "SEGMENTED_ROWS_V1"
        assert len(workbook["_SYSTEM_META"]["A2"].value) < 32_767
        assert not workbook._external_links
        assert all(
            cell.data_type != "f"
            for row in workbook[sheet_name].iter_rows()
            for cell in row
        )
        if factory_id == "huakang-b":
            assert workbook[sheet_name]["AW3"].value.startswith("白班")
            assert workbook[sheet_name]["E4"].data_type == "s"
            assert workbook[sheet_name]["AE4"].is_date
            assert workbook[sheet_name]["AF4"].is_date
        else:
            assert workbook[sheet_name]["I4"].value == "000123"
            assert workbook[sheet_name]["J4"].value == "0045"
            assert workbook[sheet_name]["I4"].data_type == "s"
            assert workbook[sheet_name]["AG4"].is_date
            assert workbook[sheet_name]["AH4"].is_date
        workbook.close()

        export_service = importlib.import_module(
            "app.services.injection_scheduling_export"
        )
        db_module = importlib.import_module("app.db")
        with db_module.SessionLocal() as db:
            manifest = export_service.verify_signed_system_meta(
                db,
                exported.content,
                factory_id=factory_id,
            )
        assert manifest["plan_id"] == plan["id"]
        assert manifest["profile_id"] == profile_id
        assert manifest["task_count"] == len(plan["tasks"])
        assert manifest["signed_row_manifest_digest"]
        assert manifest["rows"][0]["task_id"] == plan["tasks"][0]["id"]
        assert manifest["rows"][0]["completed_at_export"] >= 0
        replay = client.post(
            f"/api/injection-scheduling/plans/{plan['id']}/exports",
            json=export_payload,
        )
        assert replay.status_code == 200, replay.text
        assert replay.headers["x-idempotent-replay"] == "true"
        assert replay.content == exported.content

        round_trip = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": factory_id, "expected_revision": "0"},
            files={
                "file": (
                    f"{prefix}-round-trip.xlsx",
                    exported.content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": f"{prefix}-round-trip-preview-0001"},
        )
        assert round_trip.status_code == 201, round_trip.text
        round_trip_batch = round_trip.json()
        assert round_trip_batch["profile"]["profile_id"] == profile_id
        assert round_trip_batch["summary"]["scheduled_baseline_count"] == len(
            plan["tasks"]
        ), json.dumps(round_trip_batch, ensure_ascii=False)

        export_models = importlib.import_module(
            "app.models.injection_scheduling_export"
        )
        import_models = importlib.import_module(
            "app.models.injection_scheduling_import"
        )
        with db_module.SessionLocal() as db:
            audit = db.scalar(
                select(export_models.InjectionSchedulingExportAudit).where(
                    export_models.InjectionSchedulingExportAudit.id
                    == exported.headers["x-export-audit-id"]
                )
            )
            imported = db.get(
                import_models.InjectionSchedulingImportBatch,
                round_trip_batch["id"],
            )
            normalized = json.loads(imported.normalized_json)
        assert audit is not None
        assert audit.file_sha256 == exported.headers["x-export-sha256"]
        assert audit.plan_revision == plan["revision"]
        assert normalized["system_meta"]["audit_id"] == audit.id

        if factory_id == "huaxing":
            edited_book = load_workbook(BytesIO(exported.content))
            exported_row = manifest["rows"][0]
            edited_book[sheet_name][f"M{exported_row['source_row']}"] = (
                exported_row["completed_at_export"] + 5
            )
            edited_output = BytesIO()
            edited_book.save(edited_output)
            edited_book.close()
            edited_preview = client.post(
                "/api/injection-scheduling/imports/preview",
                data={"factory_id": factory_id, "expected_revision": "0"},
                files={
                    "file": (
                        "edited-progress.xlsx",
                        edited_output.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
                headers={"x-request-id": "phase5-edited-progress-preview-0001"},
            )
            assert edited_preview.status_code == 201, edited_preview.text
            edited_batch = edited_preview.json()
            assert "UPDATE_PROGRESS" in {
                item["action_type"]
                for item in edited_batch["reconciliation_actions"]
            }
            edited_confirm = client.post(
                f"/api/injection-scheduling/imports/{edited_batch['id']}/confirm",
                json={
                    "factory_id": factory_id,
                    "expected_revision": edited_batch["revision"],
                    "expected_plan_revision": edited_batch["plan_context"][
                        "target_draft_plan_revision"
                    ],
                    "request_id": "phase5-edited-progress-confirm-0001",
                    "confirm_mode": "merge_draft",
                    "business_date": business_date,
                    "acknowledged_blocking_issue_ids": [],
                    "expected_action_fingerprint": edited_batch[
                        "action_fingerprint"
                    ],
                },
            )
            assert edited_confirm.status_code == 200, edited_confirm.text
            refreshed = client.get(
                "/api/injection-scheduling/plans/current",
                params={"factory_id": factory_id},
            ).json()["plan"]
            assert refreshed["orders"][0]["completed_quantity"] == (
                exported_row["completed_at_export"] + 5
            )

        tampered_book = load_workbook(BytesIO(exported.content))
        tampered_book["_SYSTEM_META"]["A3"] = "0" * 64
        tampered_output = BytesIO()
        tampered_book.save(tampered_output)
        tampered_book.close()
        tampered = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": factory_id, "expected_revision": "0"},
            files={
                "file": (
                    "tampered.xlsx",
                    tampered_output.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": f"{prefix}-tampered-preview-0001"},
        )
        assert tampered.status_code == 409, tampered.text
        assert tampered.json()["detail"]["code"] == "SYSTEM_META_SIGNATURE_INVALID"

        forged_book = load_workbook(BytesIO(exported.content))
        forged_manifest = _segmented_manifest(forged_book)
        forged_manifest["audit_id"] = "isexport-forged"
        forged_manifest["export_id"] = "isexport-forged"
        forged_json = json.dumps(
            forged_manifest,
            ensure_ascii=False,
            separators=(",", ":"),
            sort_keys=True,
        )
        signing_key = b"phase5-test-signing-key-with-at-least-32-bytes"
        forged_global = dict(forged_manifest)
        forged_global.pop("rows")
        forged_book["_SYSTEM_META"]["A2"] = json.dumps(
            forged_global,
            ensure_ascii=False,
            separators=(",", ":"),
            sort_keys=True,
        )
        forged_book["_SYSTEM_META"]["A3"] = hmac.new(
            signing_key,
            forged_json.encode(),
            hashlib.sha256,
        ).hexdigest()
        forged_book["_SYSTEM_META"]["A5"] = hashlib.sha256(
            forged_json.encode()
        ).hexdigest()
        forged_output = BytesIO()
        forged_book.save(forged_output)
        forged_book.close()
        forged = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": factory_id, "expected_revision": "0"},
            files={
                "file": (
                    "forged-audit.xlsx",
                    forged_output.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": f"{prefix}-forged-audit-preview-0001"},
        )
        assert forged.status_code == 409, forged.text
        assert forged.json()["detail"]["code"] == "SYSTEM_META_EXPORT_AUDIT_INVALID"

        config = importlib.import_module("app.core.config")
        config.settings.injection_scheduling_export_signing_key = (
            "phase5-new-signing-key-with-at-least-32-bytes"
        )
        config.settings.injection_scheduling_export_signing_key_id = "phase5-test-v2"
        config.settings.injection_scheduling_export_verification_keys_json = json.dumps(
            {
                "phase5-test-v1": (
                    "phase5-test-signing-key-with-at-least-32-bytes"
                )
            }
        )
        rotated = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": factory_id, "expected_revision": "0"},
            files={
                "file": (
                    "historical-key.xlsx",
                    exported.content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": f"{prefix}-historical-key-preview-0001"},
        )
        assert rotated.status_code == 201, rotated.text
        config.settings.injection_scheduling_export_verification_keys_json = "{}"
        revoked = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": factory_id, "expected_revision": "0"},
            files={
                "file": (
                    "revoked-key.xlsx",
                    exported.content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": f"{prefix}-revoked-key-preview-0001"},
        )
        assert revoked.status_code == 409, revoked.text
        assert revoked.json()["detail"]["code"] == (
            "SYSTEM_META_SIGNING_KEY_REVOKED"
        )

        profile_models = importlib.import_module(
            "app.models.injection_scheduling_import"
        )
        with db_module.SessionLocal() as db:
            profile_record = db.get(
                profile_models.InjectionSchedulingImportProfile,
                profile_id,
            )
            assert profile_record is not None
            profile_record.status = "RETIRED"
            profile_record.lifecycle_revision += 1
            db.commit()
        config.settings.injection_scheduling_export_verification_keys_json = json.dumps(
            {
                "phase5-test-v1": (
                    "phase5-test-signing-key-with-at-least-32-bytes"
                )
            }
        )
        retired_signed = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": factory_id, "expected_revision": "0"},
            files={
                "file": (
                    "retired-signed.xlsx",
                    exported.content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": f"{prefix}-retired-signed-preview-0001"},
        )
        assert retired_signed.status_code == 201, retired_signed.text
        assert retired_signed.json()["profile"]["profile_id"] == profile_id
        unsigned_retired = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": factory_id, "expected_revision": "0"},
            files={
                "file": (
                    "retired-unsigned.xlsx",
                    source,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": f"{prefix}-retired-unsigned-preview-0001"},
        )
        assert unsigned_retired.status_code == 201, unsigned_retired.text
        assert unsigned_retired.json()["batch_state"] == "MAPPING_REQUIRED"


def test_phase5_system_standard_export_is_explicit_and_has_no_unsafe_key_default(
    monkeypatch,
):
    monkeypatch.setenv(
        "INJECTION_SCHEDULING_EXPORT_SIGNING_KEY",
        "phase5-test-signing-key-with-at-least-32-bytes",
    )
    monkeypatch.setenv("INJECTION_SCHEDULING_EXPORT_SIGNING_KEY_ID", "phase5-test-v1")
    with make_client(monkeypatch) as client:
        login(client, "admin", ADMIN_TEST_PASSWORD)
        _batch, plan = _take_over_plan(
            client,
            factory_id="huaxing",
            source=build_workbook(),
            prefix="phase5-standard",
            business_date="2026-08-01",
        )
        mismatch = client.post(
            f"/api/injection-scheduling/plans/{plan['id']}/exports",
            json={
                "factory_id": "huaxing",
                "expected_plan_revision": plan["revision"],
                "request_id": "phase5-profile-mismatch-0001",
                "export_mode": "SOURCE_COMPATIBLE",
                "profile_id": "isprofile-huakang-b-daily-v1",
                "profile_revision": 1,
            },
        )
        assert mismatch.status_code == 409
        assert mismatch.json()["detail"]["code"] == "PLAN_EXPORT_PROFILE_MISMATCH"

        exported = client.post(
            f"/api/injection-scheduling/plans/{plan['id']}/exports",
            json={
                "factory_id": "huaxing",
                "expected_plan_revision": plan["revision"],
                "request_id": "phase5-system-standard-export-0001",
                "export_mode": "SYSTEM_STANDARD",
            },
        )
        assert exported.status_code == 200, exported.text
        assert exported.headers["x-export-profile-id"] == (
            "isprofile-system-standard-v1"
        )
        round_trip = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "expected_revision": "0"},
            files={
                "file": (
                    "system-standard.xlsx",
                    exported.content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "phase5-system-standard-preview-0001"},
        )
        assert round_trip.status_code == 201, round_trip.text
        assert round_trip.json()["profile"]["profile_id"] == (
            "isprofile-system-standard-v1"
        )
        round_trip_batch = round_trip.json()
        confirmed = client.post(
            f"/api/injection-scheduling/imports/{round_trip_batch['id']}/confirm",
            json={
                "factory_id": "huaxing",
                "expected_revision": round_trip_batch["revision"],
                "expected_plan_revision": round_trip_batch["plan_context"][
                    "target_draft_plan_revision"
                ],
                "request_id": "phase5-system-standard-confirm-0001",
                "confirm_mode": "merge_draft",
                "business_date": "2026-08-01",
                "acknowledged_blocking_issue_ids": [],
                "expected_action_fingerprint": round_trip_batch[
                    "action_fingerprint"
                ],
            },
        )
        assert confirmed.status_code == 200, (
            [item["code"] for item in round_trip_batch["issues"]],
            [item["message"] for item in round_trip_batch["issues"]],
        )
        refreshed = client.get(
            "/api/injection-scheduling/plans/current",
            params={"factory_id": "huaxing"},
        ).json()["plan"]
        assert refreshed["export_profile_id"] == "isprofile-huaxing-daily-v1"
        assert len(refreshed["tasks"]) == len(plan["tasks"])

        cross_factory = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huakang-b", "expected_revision": "0"},
            files={
                "file": (
                    "cross-factory.xlsx",
                    exported.content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "phase5-cross-factory-preview-0001"},
        )
        assert cross_factory.status_code == 409, cross_factory.text
        assert cross_factory.json()["detail"]["code"] == (
            "SYSTEM_META_SCOPE_MISMATCH"
        )

        stale = client.post(
            f"/api/injection-scheduling/plans/{plan['id']}/exports",
            json={
                "factory_id": "huaxing",
                "expected_plan_revision": plan["revision"] + 1,
                "request_id": "phase5-stale-export-0001",
                "export_mode": "SYSTEM_STANDARD",
            },
        )
        assert stale.status_code == 409, stale.text
        assert stale.json()["detail"]["code"] == "PLAN_REVISION_STALE"

        ensure_user(
            "phase5-no-export",
            "molding_operator",
            factory_id="huaxing",
            department="molding",
        )
        client.post("/api/auth/logout")
        login(client, "phase5-no-export")
        denied = client.post(
            f"/api/injection-scheduling/plans/{plan['id']}/exports",
            json={
                "factory_id": "huaxing",
                "expected_plan_revision": plan["revision"],
                "request_id": "phase5-denied-export-0001",
                "export_mode": "SYSTEM_STANDARD",
            },
        )
        assert denied.status_code == 403, denied.text
        client.post("/api/auth/logout")
        login(client, "admin", ADMIN_TEST_PASSWORD)

        db_module = importlib.import_module("app.db")
        execution_models = importlib.import_module(
            "app.models.injection_scheduling_execution"
        )
        with db_module.SessionLocal() as db:
            order = db.scalar(
                select(execution_models.InjectionSchedulingOrder).where(
                    execution_models.InjectionSchedulingOrder.factory_id
                    == "huaxing"
                )
            )
            assert order is not None
            order.remark = '=HYPERLINK("https://invalid.example","x")'
            db.commit()
        safe_export = client.post(
            f"/api/injection-scheduling/plans/{plan['id']}/exports",
            json={
                "factory_id": "huaxing",
                "expected_plan_revision": refreshed["revision"],
                "request_id": "phase5-formula-safe-export-0001",
                "export_mode": "SYSTEM_STANDARD",
            },
        )
        assert safe_export.status_code == 200, safe_export.text
        safe_book = load_workbook(BytesIO(safe_export.content), data_only=False)
        assert safe_book["计划表"]["O2"].value.startswith("'=")
        assert safe_book["计划表"]["O2"].data_type == "s"
        safe_book.close()

        config = importlib.import_module("app.core.config")
        config.settings.injection_scheduling_export_signing_key = ""
        unavailable = client.post(
            f"/api/injection-scheduling/plans/{plan['id']}/exports",
            json={
                "factory_id": "huaxing",
                "expected_plan_revision": plan["revision"],
                "request_id": "phase5-no-signing-key-0001",
                "export_mode": "SYSTEM_STANDARD",
            },
        )
        assert unavailable.status_code == 503
        assert unavailable.json()["detail"]["code"] == (
            "EXPORT_SIGNING_KEY_NOT_CONFIGURED"
        )
