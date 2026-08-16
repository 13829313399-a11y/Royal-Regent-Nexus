from __future__ import annotations

import subprocess
from io import BytesIO
from pathlib import Path

import pytest
from app.core.config import Settings
from app.schemas.document_studio import (
    DocumentBlock,
    DocumentCellEvidence,
    DocumentPageSnapshot,
    DocumentSnapshot,
    DocumentTable,
)
from app.services import pdf_translation as local_pdf_translation
from app.services.document_studio.extractors.qwen_ocr import cloud_ocr_page_numbers
from app.services.document_studio.pipelines.pdf_translation import translate_snapshot
from app.services.document_studio.renderers import office_pdf_renderer
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader, PdfWriter


def _pdf_bytes(page_count: int = 1) -> bytes:
    output = BytesIO()
    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=612, height=792)
    writer.write(output)
    return output.getvalue()


def _snapshot(text: str = "订单 PO-001 数量 0012") -> DocumentSnapshot:
    return DocumentSnapshot(
        source_artifact_id="aiart-" + "a" * 32,
        source_sha256="b" * 64,
        page_count=1,
        pages=(
            DocumentPageSnapshot(
                page_number=1,
                width=612,
                height=792,
                extraction_route="NATIVE",
                blocks=(
                    DocumentBlock(
                        block_id="p1-b1",
                        kind="PARAGRAPH",
                        bbox=(10, 10, 300, 40),
                        raw_text=text,
                        normalized_text=text,
                        confidence=0.98,
                        source="NATIVE_TEXT",
                    ),
                ),
            ),
        ),
    )


def _table_snapshot() -> DocumentSnapshot:
    cells = (
        DocumentCellEvidence(
            cell_id="table-1-r1-c1",
            raw_text="0012",
            normalized_value="0012",
            value_type="IDENTIFIER",
            confidence=0.9,
            source_page=1,
            source_bbox=(10, 10, 60, 30),
            sources=("NATIVE_TEXT",),
        ),
        DocumentCellEvidence(
            cell_id="table-1-r1-c2",
            raw_text="42.00",
            normalized_value="42.00",
            value_type="TEXT",
            confidence=0.9,
            source_page=1,
            source_bbox=(60, 10, 120, 30),
            sources=("NATIVE_TEXT",),
        ),
    )
    return DocumentSnapshot(
        source_artifact_id="aiart-" + "a" * 32,
        source_sha256="b" * 64,
        page_count=1,
        pages=(
            DocumentPageSnapshot(
                page_number=1,
                width=612,
                height=792,
                extraction_route="NATIVE",
                tables=(
                    DocumentTable(
                        table_id="table-1",
                        page_number=1,
                        bbox=(10, 10, 120, 30),
                        row_count=1,
                        column_count=2,
                        cells=cells,
                        confidence=0.9,
                    ),
                ),
            ),
        ),
    )


def test_qwen_page_selection_is_a_small_request_time_heuristic() -> None:
    assert cloud_ocr_page_numbers(_snapshot()) == ()
    assert cloud_ocr_page_numbers(_snapshot(), force_all_pages=True) == (1,)


def test_pdf_translation_preserves_codes_and_numbers() -> None:
    def translator(values, direction):
        assert direction == "zh_to_en"
        return [
            value.replace("订单", "Order").replace("数量", "Quantity")
            for value in values
        ]

    result = translate_snapshot(
        _snapshot(),
        settings=Settings(_env_file=None),
        requested_direction="ZH_TO_EN",
        protected_tokens=(),
        translator=translator,
    )

    translated = result.pages[0].blocks[0].normalized_text
    assert "PO-001" in translated
    assert "0012" in translated
    assert translated.startswith("Order")


def test_pdf_translation_keeps_structured_values_and_protected_tokens() -> None:
    table_result = translate_snapshot(
        _table_snapshot(),
        settings=Settings(_env_file=None),
        requested_direction="EN_TO_ZH",
        protected_tokens=(),
        translator=lambda values, _direction: [f"Translated {value}" for value in values],
    )
    assert [
        cell.normalized_value for cell in table_result.pages[0].tables[0].cells
    ] == ["0012", "42.00"]

    received: list[str] = []

    def translator(values, _direction):
        received.extend(values)
        return [
            value.replace("Purchase order", "采购订单").replace("Quantity", "数量")
            for value in values
        ]

    result = translate_snapshot(
        _snapshot("Purchase order PO-001. Quantity 12."),
        settings=Settings(_env_file=None),
        requested_direction="EN_TO_ZH",
        protected_tokens=(),
        translator=translator,
    )
    assert received == ["Purchase order", ". Quantity"]
    assert result.pages[0].blocks[0].normalized_text == "采购订单 PO-001. 数量 12."


def test_local_pdf_translation_reuses_shared_pipeline(monkeypatch) -> None:
    monkeypatch.setattr(
        local_pdf_translation,
        "extract_local_snapshot",
        lambda **_kwargs: _snapshot(),
    )
    monkeypatch.setattr(
        local_pdf_translation,
        "render_pdf_translation",
        lambda **_kwargs: (b"translated-pdf", "order_translated.pdf", "application/pdf"),
    )

    result = local_pdf_translation.convert_pdf_translation(
        _pdf_bytes(),
        "order.pdf",
        settings=Settings(_env_file=None),
        requested_direction="ZH_TO_EN",
        protected_tokens=("PO-001",),
        translator=lambda values, _direction: [
            value.replace("订单", "Order").replace("数量", "Quantity")
            for value in values
        ],
    )
    assert result.content == b"translated-pdf"
    assert result.translated_unit_count == 1


def test_local_pdf_translation_rejects_more_than_80_pages() -> None:
    with pytest.raises(
        local_pdf_translation.PdfTranslationConversionError,
        match="单次最多翻译 80 页",
    ):
        local_pdf_translation.convert_pdf_translation(
            _pdf_bytes(81),
            "large.pdf",
            settings=Settings(_env_file=None),
            translator=lambda values, _direction: list(values),
        )


def test_office_renderer_uses_temporary_profile_and_validates_output(
    monkeypatch,
) -> None:
    source = BytesIO()
    document = Document()
    document.add_paragraph("受限 Office 渲染")
    document.save(source)
    monkeypatch.setattr(office_pdf_renderer.shutil, "which", lambda _value: "libreoffice")
    monkeypatch.setattr(office_pdf_renderer, "_validate_fonts", lambda _value: ())
    captured_command: list[str] = []

    def run(command, **_kwargs):
        captured_command.extend(command)
        output_dir = Path(command[command.index("--outdir") + 1])
        (output_dir / "source.pdf").write_bytes(_pdf_bytes())
        return subprocess.CompletedProcess(command, 0, b"", b"")

    monkeypatch.setattr(office_pdf_renderer.subprocess, "run", run)
    result = office_pdf_renderer.render_docx_to_pdf(
        source.getvalue(),
        "订单.docx",
        command="libreoffice",
    )

    assert result.page_count == 1
    assert len(PdfReader(BytesIO(result.content)).pages) == 1
    assert "--safe-mode" in captured_command
    assert any(value.startswith("-env:UserInstallation=file:") for value in captured_command)


def test_office_font_preflight_ignores_language_metadata(monkeypatch) -> None:
    source = BytesIO()
    document = Document()
    run = document.add_paragraph().add_run("language metadata is not a font")
    run.font.name = "Noto Sans CJK SC"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    language = OxmlElement("w:lang")
    language.set(qn("w:eastAsia"), "en-US")
    run._element.get_or_add_rPr().append(language)
    document.save(source)

    requested_fonts = office_pdf_renderer._requested_fonts(source.getvalue())
    assert "en-US" not in requested_fonts
    assert "Noto Sans CJK SC" in requested_fonts
    monkeypatch.setattr(
        office_pdf_renderer,
        "_installed_fonts",
        lambda: frozenset({"noto sans cjk sc", "dejavu sans mono"}),
    )
    assert office_pdf_renderer._validate_fonts(source.getvalue()) == ()
