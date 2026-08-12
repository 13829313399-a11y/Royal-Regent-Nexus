import asyncio
import hashlib
import json
from datetime import date
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from app.core.config import Settings
from app.services.ai.providers import FakeProvider, ProviderResponse, ProviderToolCall
from app.services.ai.workbook_inspection import (
    WorkbookInspectionError,
    inspect_workbook,
)
from app.services.ai.workbook_mapping import (
    WorkbookMappingError,
    propose_workbook_mapping,
)


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "啤货表"
    sheet.append(["模具编号", "产品名称", "数量", "客户显示编号", "计算"])
    sheet.append(["M-001", "产品A", 100, "00123", "=C2*2"])
    sheet["D2"].number_format = "00000"
    sheet.append(["M-002", "产品B", 200, date(2026, 8, 12), "=SUM(C2:C3)"])
    sheet.merge_cells("A5:B5")
    sheet.row_dimensions[6].hidden = True
    sheet.column_dimensions["F"].hidden = True
    hidden = workbook.create_sheet("内部核对")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "不应作为样例泄露的值"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_workbook_inspection_is_read_only_redacted_and_deterministic() -> None:
    content = workbook_bytes()
    before = bytes(content)
    first = inspect_workbook(
        source_file_name="safe.xlsx",
        content=content,
        factory_id="huaxing",
    )
    second = inspect_workbook(
        source_file_name="safe.xlsx",
        content=content,
        factory_id="huaxing",
    )
    assert content == before
    assert first.snapshot_sha256 == second.snapshot_sha256
    assert first.source_lineage.source_sha256 == hashlib.sha256(content).hexdigest()
    assert first.risk_level == "PREVIEW_WITH_AUDIT"
    assert first.tool_name == "workbook.inspect"
    assert first.sheet_count == 2
    assert first.hidden_sheet_count == 1
    assert first.formula_cell_count == 2
    source_sheet = first.sheets[0]
    assert source_sheet.merged_regions == ["A5:B5"]
    assert source_sheet.hidden_row_count == 1
    assert source_sheet.hidden_column_count == 1
    assert source_sheet.candidate_headers[0].cells[0].header_text == "模具编号"
    samples = {item.cell_ref: item for item in source_sheet.redacted_samples}
    assert samples["D2"].leading_zero is True
    assert samples["D2"].raw_sample == "<DIGITS len=5 leading_zero=yes>"
    assert samples["E2"].value_kind == "FORMULA"
    assert samples["E2"].formula_shape == "=C2*#"
    serialized = first.model_dump_json()
    assert "M-001" not in serialized
    assert "产品A" not in serialized
    assert "00123" not in serialized
    assert "不应作为样例泄露的值" not in serialized


@pytest.mark.parametrize(
    ("file_name", "content", "code"),
    [
        ("../unsafe.xlsx", workbook_bytes(), "WORKBOOK_UNSAFE_PATH"),
        ("legacy.xls", b"not-xls", "WORKBOOK_UNSUPPORTED_FORMAT"),
        ("fake.xlsx", b"not-ooxml", "WORKBOOK_INVALID_OOXML"),
    ],
)
def test_workbook_inspection_blocks_unsafe_inputs(
    file_name: str,
    content: bytes,
    code: str,
) -> None:
    with pytest.raises(WorkbookInspectionError) as captured:
        inspect_workbook(
            source_file_name=file_name,
            content=content,
            factory_id="huaxing",
        )
    assert captured.value.code == code


def test_workbook_inspection_blocks_zip_bomb_before_parsing() -> None:
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "x" * 1_500_000)
    with pytest.raises(WorkbookInspectionError) as captured:
        inspect_workbook(
            source_file_name="bomb.xlsx",
            content=output.getvalue(),
            factory_id="huaxing",
        )
    assert captured.value.code == "WORKBOOK_ZIP_BOMB"


def mapping_settings() -> Settings:
    return Settings(
        _env_file=None,
        ai_enabled=True,
        ai_provider="fake",
        ai_default_model="qwen3.7-plus",
        ai_cloud_workbook_mapping_enabled=True,
    )


def test_mapping_proposal_uses_existing_catalog_and_stays_profile_draft() -> None:
    snapshot = inspect_workbook(
        source_file_name="demand.xlsx",
        content=workbook_bytes(),
        factory_id="huaxing",
    )
    arguments = {
        "mappings": [
            {
                "canonical_field": "source_mold_no",
                "source_header": "模具编号",
                "source_sheet": "啤货表",
                "source_column": "A",
                "confidence": 0.98,
                "reason": "表头与现有规范字段别名完全一致",
                "transformer": "identifier",
            },
            {
                "canonical_field": "product_name",
                "source_header": "产品名称",
                "source_sheet": "啤货表",
                "source_column": "B",
                "confidence": 0.97,
                "reason": "表头与产品名称字段语义一致",
                "transformer": "text",
            },
            {
                "canonical_field": "order_quantity",
                "source_header": "数量",
                "source_sheet": "啤货表",
                "source_column": "C",
                "confidence": 0.96,
                "reason": "数量列与需求数量字段一致",
                "transformer": "number",
            },
        ],
        "warnings": ["工作簿数据只作为不可信来源处理"],
    }
    provider = FakeProvider(
        response=ProviderResponse(
            tool_calls=(
                ProviderToolCall(
                    call_id="mapping-1",
                    name="mapping.submit_proposal",
                    arguments_json=json.dumps(arguments, ensure_ascii=False),
                ),
            ),
            response_id="mapping-response",
        )
    )
    result = asyncio.run(
        propose_workbook_mapping(
            snapshot=snapshot,
            document_kind="DEMAND_ORDER",
            provider=provider,
            settings=mapping_settings(),
            request_id="request-mapping-1",
        )
    )
    assert result.profile_registry == "injection_scheduling_import_profiles"
    assert result.target_status == "PROFILE_DRAFT"
    assert result.approval_required is True
    assert result.known_profile_id == "isprofile-demand-order-shared-v1"
    assert result.missing_required_fields == []
    assert "Task" in " ".join(result.warnings)
    assert result.source_sha256 == snapshot.source_lineage.source_sha256
    request_payload = provider.requests[0].input[1].content
    assert isinstance(request_payload, str)
    assert "M-001" not in request_payload
    assert "产品A" not in request_payload


def test_mapping_proposal_rejects_invented_source_or_transformer() -> None:
    snapshot = inspect_workbook(
        source_file_name="demand.xlsx",
        content=workbook_bytes(),
        factory_id="huaxing",
    )
    arguments = {
        "mappings": [
            {
                "canonical_field": "source_mold_no",
                "source_header": "不存在的表头",
                "source_sheet": "啤货表",
                "source_column": "Z",
                "confidence": 1,
                "reason": "模型尝试凭空构造一个来源字段",
                "transformer": "identifier",
            }
        ],
        "warnings": [],
    }
    provider = FakeProvider(
        response=ProviderResponse(
            tool_calls=(
                ProviderToolCall(
                    call_id="mapping-invalid",
                    name="mapping.submit_proposal",
                    arguments_json=json.dumps(arguments, ensure_ascii=False),
                ),
            )
        )
    )
    with pytest.raises(WorkbookMappingError) as captured:
        asyncio.run(
            propose_workbook_mapping(
                snapshot=snapshot,
                document_kind="DEMAND_ORDER",
                provider=provider,
                settings=mapping_settings(),
                request_id="request-mapping-invalid",
            )
        )
    assert captured.value.code == "AI_MAPPING_INVALID_OUTPUT"
