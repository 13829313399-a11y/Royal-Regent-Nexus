import importlib
from io import BytesIO
from pathlib import Path
import json

import pytest
from openpyxl import Workbook, load_workbook

from test_internal_quote_api import make_client, login, logout
from test_internal_quote_alternatives import create, fill, fork


def body(quotes):
    return {"versions": [{"quote_id": row["id"], "revision": row["header_revision"]} for row in quotes]}


def exports(client, quote):
    return client.get(f"/api/internal-quotes/{quote['id']}/exports").json()


def test_summary_handles_component_totals_safe_text_and_rejects_unreadable_prices():
    from app.services.internal_quote_alternative_export import combine_alternative_workbooks
    from fastapi import HTTPException
    book = Workbook()
    book.active.title = '报价明细'
    calculation = book.create_sheet('_报价明细计算')
    calculation['B5'] = '报价合计'; calculation['D5'] = '=SUM(10,20)'; calculation['R6'] = 3000
    data = book.create_sheet('结构化数据')
    data.append([]); data.append([]); data.append([])
    data.append(['payload', 'sales', '', '', '', '', '', '', 1, 1,
                 json.dumps({'packaging_materials': [{'item': '=HYPERLINK("external")', 'specification': '安全文本'}]})])
    def item():
        buffer = BytesIO(); book.save(buffer)
        return dict(label='=方案/[]'*8, scenario='=危险公式', version='A-V1', content=buffer.getvalue(),
                    exported_at='2026-10-15', file_name='原始.xlsx', sha256='sample')
    result = load_workbook(BytesIO(combine_alternative_workbooks([item(), item()])))
    sheet = result['方案汇总']
    assert sheet['A5'].data_type == sheet['C5'].data_type == 's'
    assert sheet['F5'].value == 30 and sheet['E5'].value == 3000
    assert len(result.sheetnames) == len(set(result.sheetnames)) and all(len(n) <= 31 for n in result.sheetnames)
    result.close()
    calculation['D5'] = '=UNSUPPORTED(1)'
    with pytest.raises(HTTPException, match='复制新版本'):
        combine_alternative_workbooks([item(), item()])
    book.close()


def test_comparison_uses_frozen_prices_packaging_and_preserves_original_bytes(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "alternative-export", "admin", "*", "*")
        original = fill(client, create(client, "COMPARE"))
        second = fork(client, original, name="开窗彩盒/透明内托")
        sales = next(row for row in second["sections"] if row["department"] == "sales")
        sales["payload"]["packaging_materials"][0].update(item="开窗彩盒", specification="透明内托", unit_price_rmb=2)
        sales["payload"]["shipping"]["markup_tiers"] = [
            {"moq": 3000, "markup_x": 1.25, "include_in_output": True},
            {"moq": 5000, "markup_x": 1.2, "include_in_output": True},
            {"moq": 10000, "markup_x": 1.15, "include_in_output": False},
        ]
        saved = client.put(f"/api/internal-quotes/{second['id']}/sections/sales", json={
            "revision": sales["revision"], "payload": sales["payload"], "reason": "包装方案测试"})
        assert saved.status_code == 200, saved.text
        second = client.get(f"/api/internal-quotes/{second['id']}").json()
        first_export = client.post(f"/api/internal-quotes/{original['id']}/direct-issue", json={"revision": original["header_revision"]}).json()
        frozen_url = f"/api/internal-quotes/{original['id']}/exports/{first_export['id']}/download"
        frozen = client.get(frozen_url).content
        endpoint = f"/api/internal-quotes/{original['id']}/alternatives-export"
        response = client.post(endpoint, json=body([original, second]))
        assert response.status_code == 200, response.text[:1000]
        assert client.get(frozen_url).content == frozen
        workbook = load_workbook(BytesIO(response.content))
        visible = [sheet for sheet in workbook if sheet.sheet_state == "visible"]
        assert len(visible) == 3 and visible[0].title == "方案汇总"
        summary = visible[0]
        assert summary.max_row == 7
        assert [summary.cell(row, 5).value for row in (5, 6, 7)] == [5000, 3000, 5000]
        assert "开窗彩盒" in summary["C6"].value and "透明内托" in summary["C6"].value
        assert summary["F6"].value > summary["F5"].value
        assert summary["G5"].hyperlink and summary.freeze_panes == "C5"
        # Every summary amount must equal the exact source workbook cell, not a new server calculation.
        view = importlib.import_module("app.services.internal_quote_workbook_view")
        values = view.WorkbookValues(workbook)
        for row in (5, 6, 7):
            source = summary.cell(row, 6).comment.text.split("来源：")[1].split("\n")[0]
            sheet, address = source.rsplit("!", 1)
            assert summary.cell(row, 6).value == pytest.approx(values.cell(sheet, address))
        workbook.close()
        artifact = Path(__file__).parents[1] / ".pytest-tmp" / "alternative-export-acceptance.xlsx"
        artifact.write_bytes(response.content)
        # Formula/layout code changes cannot regenerate either frozen version on retry.
        artifacts = importlib.import_module("app.services.internal_quote_artifacts")
        monkeypatch.setattr(artifacts, "build_internal_quote_workbook", lambda *a, **k: pytest.fail("regenerated frozen export"))
        repeated = client.post(endpoint, json=body([original, second]))
        assert repeated.status_code == 200, repeated.text[:1000]
        assert len(exports(client, original)) == len(exports(client, second)) == 1
        assert client.get(frozen_url).content == frozen


def test_merge_failure_rolls_back_issues_and_handoffs(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "alternative-rollback", "admin", "*", "*")
        source = fill(client, create(client, "ROLLBACK"))
        target = fork(client, source)
        endpoint = f"/api/internal-quotes/{source['id']}/alternatives-export"
        service = importlib.import_module("app.services.internal_quote_alternative_export")
        with monkeypatch.context() as failure:
            failure.setattr(service, "combine_alternative_workbooks", lambda _: (_ for _ in ()).throw(ValueError("workbook failed")))
            with pytest.raises(ValueError, match="workbook failed"):
                client.post(endpoint, json=body([source, target]))
        models = importlib.import_module("app.models.internal_quote")
        db_module = importlib.import_module("app.db")
        with db_module.SessionLocal() as db:
            assert db.query(models.InternalQuoteArtifactHandoff).count() == 0
        for quote in (source, target):
            current = client.get(f"/api/internal-quotes/{quote['id']}").json()
            assert current["header_revision"] == quote["header_revision"]
            assert current["final_release_status"] != "issued" and exports(client, quote) == []


def test_validates_membership_revisions_archive_permissions_and_incomplete_versions(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "alternative-boundaries", "admin", "*", "*")
        source = create(client, "BOUNDARIES")
        incomplete = fork(client, source)
        source = fill(client, source)
        other = create(client, "UNRELATED")
        endpoint = f"/api/internal-quotes/{source['id']}/alternatives-export"
        assert client.post(endpoint, json=body([source])).status_code == 422
        for selected in ([source, source], [source, other]):
            assert client.post(endpoint, json=body(selected)).status_code == 400
        stale = body([source, incomplete]); stale["versions"][0]["revision"] += 1
        assert client.post(endpoint, json=stale).status_code == 409
        assert client.post(endpoint, json=body([source, incomplete])).status_code == 409
        assert exports(client, source) == []
        family = client.get(f"/api/internal-quotes/{source['id']}/alternatives").json()
        archived = client.post(f"/api/internal-quotes/{incomplete['id']}/alternative-archive", json={
            "family_revision": family["revision"], "archived": True, "reason": "暂不采用"})
        assert archived.status_code == 200, archived.text
        assert client.post(endpoint, json=body([source, incomplete])).status_code == 409
        logout(client); login(client, "alternative-no-export", "engineer", "engineering")
        assert client.post(endpoint, json=body([source, incomplete])).status_code == 403
        logout(client); login(client, "alternative-other-factory", "sales_customer_supervisor", "sales-business", "huadeng")
        assert client.post(endpoint, json=body([source, incomplete])).status_code == 403


def test_legacy_final_export_is_reused_without_current_formula_validation(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "alternative-legacy", "admin", "*", "*")
        source = fill(client, create(client, "LEGACY"))
        target = fork(client, source)
        issued = client.post(f"/api/internal-quotes/{source['id']}/direct-issue", json={"revision": source["header_revision"]}).json()
        db_module = importlib.import_module("app.db")
        models = importlib.import_module("app.models.internal_quote")
        with db_module.SessionLocal() as db:
            quote = db.get(models.InternalQuote, source["id"])
            quote.module_version = "v3"; quote.final_release_status = "approved"; quote.formula_version = "historical"
            db.get(models.InternalQuoteExportFile, issued["id"]).release_stage = "p4_final_approved"
            db.commit()
        frozen_url = f"/api/internal-quotes/{source['id']}/exports/{issued['id']}/download"
        before = client.get(frozen_url).content
        response = client.post(f"/api/internal-quotes/{source['id']}/alternatives-export", json=body([source, target]))
        assert response.status_code == 200, response.text[:1000]
        assert client.get(frozen_url).content == before and len(exports(client, source)) == 1
        with db_module.SessionLocal() as db:
            record = db.get(models.InternalQuoteExportFile, issued['id'])
            manifest = json.loads(record.export_manifest_json)
            manifest['final_release_revision'] = -1
            record.export_manifest_json = json.dumps(manifest)
            db.commit()
        missing = client.post(f"/api/internal-quotes/{source['id']}/alternatives-export", json=body([source, target]))
        assert missing.status_code == 409 and '不能重新生成历史报价' in missing.text
        assert len(exports(client, source)) == 1
