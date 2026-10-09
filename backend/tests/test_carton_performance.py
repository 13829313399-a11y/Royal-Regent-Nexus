"""Synthetic large-data and lifecycle contracts; no production data required."""
import importlib
import pickle
import time
from io import BytesIO
from threading import Event

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook


def workbook(rows, sheets=1):
    book = Workbook()
    sheet = book.active
    sheet.append(["合同号", "货号", "数量", "客名", "产品名称"])
    for index in range(rows):
        sheet.append([f"SC-{index}", f"ITEM-{index}", 100, "TEST", "Synthetic"])
    for index in range(1, sheets):
        book.copy_worksheet(sheet).title = f"Sheet{index}"
    output = BytesIO(); book.save(output); book.close()
    return output.getvalue()


def test_parse_complete_5000_and_reject_entire_oversized_batch():
    from app.services.carton_procurement_imports import parse_carton_file
    parsed = parse_carton_file("WEEKLY_SCHEDULE", "large.xlsx", workbook(5000))
    assert len(parsed["rows"]) == 5000
    assert parsed["rows"][-1]["item_no"] == "ITEM-4999"
    with pytest.raises(HTTPException, match="5000"):
        parse_carton_file("WEEKLY_SCHEDULE", "large.xlsx", workbook(5001))
    with pytest.raises(HTTPException, match="5000"):
        parse_carton_file("WEEKLY_SCHEDULE", "large.xlsx", workbook(3000, 2))


def test_real_worker_parses_without_parent_database():
    from app.services.carton_file_jobs import run_file_job
    parsed = run_file_job("parse", "WEEKLY_SCHEDULE", "large.xlsx", workbook(601))
    assert len(parsed["rows"]) == 601


def test_task_owner_factory_capacity_retry_and_completion(monkeypatch):
    jobs = importlib.import_module("app.services.carton_file_jobs")
    gate = Event()
    def compute(folder, operation, args):
        assert gate.wait(5)
        (folder / "input.pkl").write_bytes(pickle.dumps((operation, args)))
        (folder / "result.pkl").write_bytes(pickle.dumps({"value": {"rows": [1]}}))
    monkeypatch.setattr(jobs, "_compute", compute)
    def ready(identifier):
        deadline = time.monotonic() + 5
        job = jobs.get_job(identifier, "alice", "huakang-b")
        while job.status == "PROCESSING" and time.monotonic() < deadline:
            time.sleep(.01)
        assert job.status == "READY"
        return job
    try:
        first = jobs.start_job("alice", "huakang-b", "parse", ("same",), {})
        assert jobs.start_job("alice", "huakang-b", "parse", ("same",), {})["id"] == first["id"]
        with pytest.raises(HTTPException) as busy:
            jobs.start_job("alice", "huakang-b", "parse", ("different",), {})
        assert busy.value.status_code == 429
        for owner, factory in [("bob", "huakang-b"), ("alice", "huaxing")]:
            with pytest.raises(HTTPException) as forbidden:
                jobs.get_job(first["id"], owner, factory)
            assert forbidden.value.status_code == 404
        gate.set()
        job = ready(first["id"])
        with jobs.result(job) as (value, args):
            assert value == {"rows": [1]} and args == ("same",)
            job.status, job.batch_id = "COMPLETED", "durable-batch"
        # Completed results remain retryable, but a new upload after undo gets a new job.
        with jobs.result(job):
            assert job.batch_id == "durable-batch"
        for index in range(12):
            fresh = jobs.start_job("alice", "huakang-b", "parse", ("same",), {})
            assert fresh["id"] != first["id"]
            ready(fresh["id"]).status = "COMPLETED"
        assert len(jobs._jobs) <= jobs.MAX_JOBS
    finally:
        gate.set()
        for identifier in list(jobs._jobs):
            jobs._expire(identifier)


def test_projection_matches_legacy_with_bounded_sql_and_indexed_matching():
    import pkgutil
    import app.models
    from sqlalchemy import create_engine, event, insert, select
    from sqlalchemy.orm import Session
    for module in pkgutil.iter_modules(app.models.__path__):
        importlib.import_module('app.models.' + module.name)
    from app.db import Base
    from app.models.carton_procurement import CartonOrder, CartonOrderLine
    from app.services.carton_procurement import order_out, list_orders
    from app.services.carton_order_projection import order_page_out
    from app.services.carton_procurement_imports import _match_rows
    from app.services.carton_procurement_export import build_combined_purchase_order_workbook
    from datetime import datetime
    engine = create_engine("sqlite://")
    Base.metadata.create_all(engine)
    queries = []
    event.listen(engine, "before_cursor_execute", lambda *args: queries.append(args[2]))
    audit = dict(created_by="test", updated_by="test", created_at="2026-09-30", updated_at="2026-09-30")
    with Session(engine) as db:
        db.execute(insert(CartonOrder), [dict(id=f"O{i}", factory_id="huakang-b", order_no=f"CT{i:05}",
            customer_code="TEST", customer_name="Synthetic", supplier_id="S", supplier_name_snapshot="Synthetic", contract_no=f"SC-{i}", item_no=f"ITEM-{i}",
            product_order_quantity=100, order_date="2026-09-30", due_date="2026-10-30", status="CONFIRMED", **audit) for i in range(5000)])
        db.execute(insert(CartonOrderLine), [dict(id=f"L{i}", factory_id="huakang-b", order_id=f"O{i}",
            line_no=1, customer_code="TEST", contract_no=f"SC-{i}", item_no=f"ITEM-{i}", packaging_type="BOX",
            paper_quality="A", specification="30*20*10", usage_quantity=1, required_quantity=100,
            unit="PCS", unit_price=1, currency="HKD") for i in range(5000)])
        db.commit()
        _, page = list_orders(db, "huakang-b", limit=200)
        expected = [order_out(db, order) for order in page[:3]]
        queries.clear(); started = time.perf_counter()
        projected = order_page_out(db, page)
        elapsed = time.perf_counter() - started
        query_count = len(queries)
        assert projected[:3] == expected
        assert query_count <= 40, query_count
        queries.clear(); order_page_out(db, page[:1])
        assert len(queries) == query_count
        rows = [dict(contract_no=f"SC-{i}", item_no=f"ITEM-{i}", quantity=100) for i in range(5000)]
        started = time.perf_counter(); _match_rows(db, "huakang-b", "WEEKLY_SCHEDULE", rows)
        matched_seconds = time.perf_counter() - started
        assert all(row["match_status"] == "MATCHED" and row["order_id"] == f"O{i}" for i, row in enumerate(rows))
        assert matched_seconds < 15  # Before the index this workload took about 75 seconds.
        stats = {}
        total, filtered = list_orders(db, "huakang-b", search="ITEM-4999", statistics=stats)
        assert total == 1 and filtered[0].id == "O4999" and stats["pending"] == 1
        assert list_orders(db, "huakang-b", search="待下单")[0] == 5000
        assert list_orders(db, "huakang-b", search="Synthetic SC-4999")[0] == 1
        assert list_orders(db, "huakang-b", search="%")[0] == 0
        assert list_orders(db, "huakang-b", search="_")[0] == 0
        lines = {line.order_id: line for line in db.scalars(select(CartonOrderLine))}
        contents = build_combined_purchase_order_workbook([(order, [lines[order.id]]) for order in page[:100]], generated_at=datetime(2026, 10, 1))
        book = load_workbook(BytesIO(contents)); assert book.active.max_row >= 100
        book.close()
        from types import SimpleNamespace
        large = [(order, [SimpleNamespace(**{**{column.name: getattr(lines[order.id], column.name)
            for column in CartonOrderLine.__table__.columns}, "line_no": number}) for number in range(1, 51)]) for order in page[:100]]
        started = time.perf_counter()
        contents = build_combined_purchase_order_workbook(large, generated_at=datetime(2026, 10, 1))
        export_seconds = time.perf_counter() - started
        book = load_workbook(BytesIO(contents))
        assert book.active["J5007"].value == "BOX"
        assert book.active["N8"].value == "=ROUNDUP(I8/M8,0)"
        assert book.active["A8"].border.top.style == "medium"
        book.close()
        print(f"PERFORMANCE page200_sql={query_count} page_seconds={elapsed:.3f} match5000_seconds={matched_seconds:.3f}")
        print(f"PERFORMANCE export100_orders_5000_lines_seconds={export_seconds:.3f} bytes={len(contents)}")
    engine.dispose()


def test_job_api_stays_responsive_completes_once_and_preserves_projection(monkeypatch):
    from test_molding_sample_api import make_client, login_as
    from test_carton_procurement_api import _create_order, _freeze_carton_time
    from test_carton_receipt_correction_api import receipt
    from sqlalchemy import select
    from fastapi.encoders import jsonable_encoder
    base = "/api/carton-procurement"
    with make_client(monkeypatch) as client:
        login_as(client, "admin"); _freeze_carton_time(monkeypatch)
        order = _create_order(client)
        receipt(client, order, post=False)
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonOrder
        from app.services.carton_procurement import order_out
        with SessionLocal() as db:
            expected = jsonable_encoder(order_out(db, db.scalar(select(CartonOrder).where(CartonOrder.id == order["id"]))))
        listed = client.get(base + "/orders", params={"factory_id": "huaxing"})
        assert listed.status_code == 200, listed.text
        assert listed.json()["items"] == [expected]
        selected = client.post(base + "/orders/selection", json={"factory_id": "huaxing", "order_nos": [order["order_no"]]})
        assert selected.status_code == 200, selected.text
        assert selected.json()["items"] == [expected]
        index = client.get(base + "/schedule-orders", params={"factory_id": "huaxing"})
        assert index.status_code == 200 and index.json()["total"] == 1
        assert "lines" not in index.json()["items"][0]
        assert client.post(base + "/orders/selection", json={"factory_id": "huakang-b", "order_nos": [order["order_no"]]}).status_code == 403
        login_as(client, "admin")
        assert client.post(base + "/orders/selection", json={"factory_id": "huakang-b", "order_nos": [order["order_no"]]}).json()["items"] == []
        jobs = importlib.import_module("app.services.carton_file_jobs")
        original = jobs._compute; gate = Event()
        def waiting_compute(*args):
            assert gate.wait(5)
            return original(*args)
        monkeypatch.setattr(jobs, "_compute", waiting_compute)
        response = client.post(base + "/file-jobs/imports", params={"factory_id": "huaxing", "import_type": "WEEKLY_SCHEDULE"},
            data={"customer_code": "DICKIE"}, files={"file": ("synthetic.xlsx", workbook(601))})
        try:
            assert response.status_code == 202, response.text
            identifier = response.json()["id"]
            assert client.get("/health").status_code == 200
            assert client.get(base + "/orders", params={"factory_id": "huaxing"}).status_code == 200
            assert client.post(base + f"/file-jobs/{identifier}/complete", params={"factory_id": "huaxing"}).status_code == 409
            gate.set(); deadline = time.monotonic() + 10
            while time.monotonic() < deadline:
                status = client.get(base + f"/file-jobs/{identifier}", params={"factory_id": "huaxing"}).json()
                if status["status"] != "PROCESSING": break
                time.sleep(.05)
            assert status["status"] == "READY", status
            done = client.post(base + f"/file-jobs/{identifier}/complete", params={"factory_id": "huaxing"})
            assert done.status_code == 200, done.text
            assert done.json()["parse_summary"]["row_count"] == 601
            assert client.post(base + f"/file-jobs/{identifier}/complete", params={"factory_id": "huaxing"}).json()["id"] == done.json()["id"]
        finally:
            gate.set()
            for identifier in list(jobs._jobs): jobs._expire(identifier)
