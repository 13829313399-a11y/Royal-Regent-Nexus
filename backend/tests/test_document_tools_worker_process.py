"""Opt-in real worker-process smoke with an exclusively isolated database.

RR_DOCUMENT_WORKER_SMOKE=1 python -m pytest tests/test_document_tools_worker_process.py
No HTTP server, business database, external OCR, or production service is used.
"""
import hashlib
import json
import os
import shutil
import subprocess
import sys
import time
import uuid
from pathlib import Path

import pytest


@pytest.mark.skipif(os.getenv("RR_DOCUMENT_WORKER_SMOKE") != "1", reason="explicit real Office/worker-process smoke opt-in")
def test_real_worker_process_all_seven_directions():
    from docx import Document
    from docx.shared import Inches
    from openpyxl import Workbook, load_workbook
    from PIL import Image
    from pypdf import PdfReader
    from sqlalchemy import create_engine, event, inspect, select, text
    from sqlalchemy.orm import sessionmaker

    from app.db import Base
    from app.models.auth import AuthUser
    from app.models.document_tools import (
        DocumentToolArtifact as Artifact,
        DocumentToolCorrection as Correction,
        DocumentToolJob as Job,
        DocumentToolSource as Source,
    )

    repo = Path(__file__).resolve().parents[2]
    output_base = Path(os.environ.get("RR_DOCUMENT_WORKER_SMOKE_DIR", "D:/RR-test-artifacts/document-worker-smoke")).resolve()
    # Refuse the actual business data subtree even if a caller changes opt-in settings.
    assert not output_base.is_relative_to((repo / "backend/data").resolve())
    run_dir = output_base / (time.strftime("%Y%m%d-%H%M%S") + "-" + uuid.uuid4().hex[:8])
    run_dir.mkdir(parents=True, exist_ok=False)
    store = run_dir / "storage"
    inputs = run_dir / "inputs"
    store.mkdir(); inputs.mkdir()
    database = run_dir / "isolated.sqlite"
    database_url = "sqlite:///" + database.as_posix()
    child_env = os.environ.copy()
    child_env.update({
        "DATABASE_URL": database_url,
        "DOCUMENT_TOOLS_STORAGE_DIR": str(store),
        "DOCUMENT_TOOLS_AI_MODE": "off",
        "DOCUMENT_TOOLS_QWEN_API_KEY": "",
        "DOCUMENT_TOOLS_QWEN_BASE_URL": "",
        "DOCUMENT_TOOLS_ENABLED": "true",
        "DOCUMENT_TOOLS_WORKER_CONCURRENCY": "1",
        "SEED_DEFAULT_ACCOUNTS": "false",
        "APP_ENV": "test",
        "AUTHZ_MODE": "legacy",
        "AUTHZ_WRITES_ENABLED": "false",
        "PYTHONIOENCODING": "utf-8",
        "PYTHONPATH": str(repo / "backend"),
    })
    assert child_env["DATABASE_URL"].endswith("isolated.sqlite")
    assert Path(child_env["DOCUMENT_TOOLS_STORAGE_DIR"]).is_relative_to(run_dir)

    engine = create_engine(database_url, connect_args={"check_same_thread": False})
    @event.listens_for(engine, "connect")
    def foreign_keys(connection, _):
        connection.execute("PRAGMA foreign_keys=ON")
    Base.metadata.create_all(engine, tables=[AuthUser.__table__, Source.__table__, Job.__table__, Artifact.__table__, Correction.__table__])
    sessions = sessionmaker(engine, expire_on_commit=False)
    owner = "worker-smoke-owner"
    with sessions() as db:
        db.add(AuthUser(id=owner, username=owner, password_salt="synthetic", password_hash="not-a-real-login"))
        db.commit()

    document = Document()
    document.add_paragraph("Synthetic worker source; item 00123, amount 12.50")
    table = document.add_table(rows=3, cols=3)
    for r, values in enumerate([["ITEM", "QTY", "AMOUNT"], ["00123", "12", "12.50"], ["00456", "24", "24.50"]]):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    logo = inputs / "synthetic-logo.png"
    Image.new("RGB", (60, 30), "#17877c").save(logo)
    document.add_picture(str(logo), width=Inches(.8))
    document.save(inputs / "source.docx")
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Synthetic"
    for row in [["ITEM", "QTY", "AMOUNT"], ["00123", 12, 12.5], ["00456", 24, 24.5]]:
        sheet.append(row)
    for row in (2, 3):
        sheet.cell(row, 1).number_format = "@"
        sheet.cell(row, 3).number_format = "0.00"
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 15
    sheet.column_dimensions["C"].width = 20
    sheet.print_area = "A1:C3"
    workbook.save(inputs / "source.xlsx")
    # Reproducible helper writes actual PDF font/text/vector objects.
    from test_document_tools_pdf_engine import make_pdf
    make_pdf(inputs / "source.pdf", table=True)

    report = {"run_directory": str(run_dir), "database_path": str(database), "storage_path": str(store),
              "ai_mode": "off", "external_ai_key_passed": False, "jobs": [], "overall_status": "running"}
    def save_report():
        (run_dir / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        (output_base / "latest-run.json").write_text(json.dumps({"run_directory": str(run_dir), "report_path": str(run_dir / "report.json")}, indent=2), encoding="utf-8")

    def enqueue(source_id, operation, kind="convert", options=None):
        job_id = uuid.uuid4().hex
        with sessions() as db:
            job = Job(id=job_id, owner_user_id=owner, source_id=source_id, operation=operation,
                      kind=kind, options_json={"ai_mode": "off", **(options or {})}, created_at=time.strftime("%Y-%m-%dT%H:%M:%SZ"))
            db.add(job)
            if kind == "inspect":
                source = db.get(Source, source_id)
                source.inspection_job_id = job_id
            db.commit()
        return job_id

    def run_child(job_id, operation):
        started = time.monotonic()
        result = subprocess.run([sys.executable, "-m", "app.workers.document_tools", "--once"],
            cwd=repo / "backend", env=child_env, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=180,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
        # Worker output contains only safe engine diagnostics; no env/request dump.
        (run_dir / f"{operation}-{job_id}.log").write_text(result.stdout + "\n" + result.stderr, encoding="utf-8")
        with sessions() as db:
            job = db.get(Job, job_id)
            artifacts = db.scalars(select(Artifact).where(Artifact.job_id == job_id)).all()
            detail = {"id": job_id, "operation": operation, "process_returncode": result.returncode,
                "elapsed_seconds": round(time.monotonic() - started, 3), "execution_status": job.execution_status,
                "quality_status": job.quality_status, "error_code": job.error_code, "error_message": job.error_message,
                "attempt": job.attempt, "lease_cleared": job.lease_token is None and job.lease_until is None, "artifacts": []}
            report["jobs"].append(detail); save_report()
            assert result.returncode == 0, detail
            assert job.execution_status == "succeeded", detail
            assert detail["lease_cleared"] and job.attempt == 1, detail
            assert artifacts, detail
            for artifact in artifacts:
                artifact_path = (store / artifact.storage_key).resolve()
                assert artifact_path.is_relative_to(store)
                data = artifact_path.read_bytes()
                assert len(data) == artifact.size
                assert hashlib.sha256(data).hexdigest() == artifact.sha256
                record = {"role": artifact.role, "format": artifact.format, "path": str(artifact_path),
                          "size": artifact.size, "sha256": artifact.sha256, "hash_verified": True}
                if artifact.format == "pdf":
                    reader = PdfReader(artifact_path)
                    assert reader.pages
                    extracted = " ".join(page.extract_text() for page in reader.pages)
                    record.update(pages=len(reader.pages), contains_identifier="00123" in extracted, contains_amount="12.50" in extracted)
                    if artifact.role in {"result", "preview"}:
                        assert "00123" in extracted, record
                        assert "12.50" in extracted, record
                elif artifact.format == "docx":
                    doc = Document(artifact_path)
                    content = " ".join([p.text for p in doc.paragraphs] + [cell.text for t in doc.tables for row in t.rows for cell in row.cells])
                    assert doc.tables and "00123" in content and "12.50" in content
                    record.update(editable_tables=len(doc.tables), contains_identifier=True, contains_amount=True)
                elif artifact.format == "xlsx":
                    wb = load_workbook(artifact_path, data_only=False)
                    values = [cell.value for ws in wb for row in ws for cell in row]
                    assert "00123" in values and (12.5 in values or "12.50" in values)
                    record.update(sheets=len(wb.sheetnames), contains_identifier=True, contains_amount=True)
                    wb.close()
                elif artifact.role == "ir":
                    ir = json.loads(data)
                    image_paths = [Path(b["style"]["image_path"]) for b in ir.get("blocks", []) if b.get("style", {}).get("image_path")]
                    assert all(p.is_file() and p.resolve().is_relative_to(store) for p in image_paths)
                    record["persistent_image_paths_verified"] = len(image_paths)
                    assert not ir.get("engine_manifest", {}).get("qwen_calls"), "AI must remain off"
                detail["artifacts"].append(record)
            assert any(a.role == "preview" or (a.role == "result" and a.format == "pdf") for a in artifacts), detail
            if operation.startswith("inspect-"):
                source = db.get(Source, job.source_id)
                assert source.inspection_status == "succeeded" and (store / source.manifest_key).is_file()
            save_report()

    try:
        sources = {}
        for kind in ("docx", "xlsx", "pdf"):
            source_id = uuid.uuid4().hex
            source_file = inputs / ("source." + kind)
            key = f"sources/{owner}/{source_id}/original.{kind}"
            stored = store / key; stored.parent.mkdir(parents=True); shutil.copyfile(source_file, stored)
            with sessions() as db:
                db.add(Source(id=source_id, owner_user_id=owner, original_name=source_file.name,
                    detected_type=kind, sha256=hashlib.sha256(stored.read_bytes()).hexdigest(), byte_size=stored.stat().st_size,
                    storage_key=key, created_at=time.strftime("%Y-%m-%dT%H:%M:%SZ")))
                db.commit()
            sources[kind] = source_id
            run_child(enqueue(source_id, "inspect", kind="inspect"), "inspect-" + kind)
        operations = [("word_to_pdf", "docx"), ("pdf_to_word", "pdf"), ("word_to_excel", "docx"),
                      ("excel_to_word", "xlsx"), ("pdf_to_excel", "pdf"), ("excel_to_pdf", "xlsx"), ("pdf_split", "pdf")]
        for operation, kind in operations:
            options = {"split_mode": "each"} if operation == "pdf_split" else {}
            run_child(enqueue(sources[kind], operation, options=options), operation)
        with engine.connect() as connection:
            assert connection.execute(text("PRAGMA quick_check")).scalar() == "ok"
            assert connection.execute(text("PRAGMA foreign_key_check")).all() == []
        report.update(overall_status="passed", conversion_directions=7, subprocess_jobs=10,
                      sqlite_tables=inspect(engine).get_table_names(), integrity="ok", foreign_key_errors=0)
        assert len(report["sqlite_tables"]) == 5
    except BaseException as exc:
        report.update(overall_status="failed", failure_type=type(exc).__name__)
        raise
    finally:
        save_report()
        engine.dispose()
