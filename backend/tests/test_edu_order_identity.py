from io import BytesIO

import openpyxl
import pytest
import xlwt

from app.services import customer_order_unified as service
from app.services import customer_order_huaxing_unified as mapping
from test_customer_order_huaxing_unified import template, data


def po_bytes(contracts, *, quantities=None, products=None, pos=None):
    book = xlwt.Workbook()
    sheet = book.add_sheet('订单')
    book.add_sheet('EDU日期格式')
    headers = ['客人', '国家', 'SO#', '合同', '产品型号', '数量', '收单日期', '走货期', '扣数PO#']
    for col, value in enumerate(headers):
        sheet.write(2, col, value)
    for i, contract in enumerate(contracts):
        values = ['Buki France', '法国', 12810117 + i, contract,
                  products[i] if products else 'F12-MS907-B3080000',
                  quantities[i] if quantities else 6360, '2026-09-29', '2026-12-07',
                  pos[i] if pos else '']
        for col, value in enumerate(values):
            sheet.write(i + 3, col, value)
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def preview_args(po, wb=None):
    return dict(customer_code='edu', factory_id='huaxing', received_date='2026-10-05',
                po_files=[('华兴代工订单.xls', po)], schedule_file_name='EDU.xlsx',
                schedule_content=data(wb or template('edu')))


def test_nine_distinct_contracts_with_repeated_item_export_all_three_sheets():
    contracts = [f'46000020{37 + i}/EP{25 + i}' for i in range(9)]
    args = preview_args(po_bytes(contracts))
    preview = service.create_unified_customer_preview(**args)
    assert preview['summary'] == dict(total=9, valid=9, warning=0, blocked=0)
    refs = [r['reference_no'] for r in preview['rows']]
    assert len(set(refs)) == 9
    output, _, _ = service.export_unified_customer_schedule(**args)
    book = openpyxl.load_workbook(BytesIO(output))
    for name, offset in [('ITEM表', 4), ('接单表', 3), ('正单评审表', 3)]:
        written = [row for row in book[name] if row[offset - 1].value in contracts]
        assert len(written) == 9
        assert {row[offset].value for row in written} == set(refs)
    retry = service.create_unified_customer_preview(**{**args, 'schedule_content': output})
    assert retry['summary']['blocked'] == 9
    assert all(any(i['code'] == 'duplicate_existing_order' for i in r['issues']) for r in retry['rows'])


@pytest.mark.parametrize('qty,code,skippable', [(6360, 'duplicate_batch_order_line', True),
                                              (6000, 'existing_quantity_conflict', False)])
def test_real_batch_duplicates_and_quantity_conflicts_remain_blocked(qty, code, skippable):
    preview = service.create_unified_customer_preview(**preview_args(po_bytes(['C1/EP1', 'C1/EP1'], quantities=[6360, qty])))
    row = preview['rows'][1]
    issue = next(i for i in row['issues'] if i['code'] == code)
    assert row['status'] == 'blocked'
    assert issue['can_skip'] is skippable


def test_same_contract_different_products_share_reference_but_different_po_do_not():
    args = preview_args(po_bytes(['C1', 'C1', 'C1'], products=['I1', 'I2', 'I1'], pos=['P1', 'P1', 'P2']))
    rows = service.create_unified_customer_preview(**args)['rows']
    assert rows[0]['reference_no'] == rows[1]['reference_no'] != rows[2]['reference_no']
    assert all(r['status'] == 'valid' for r in rows)


def test_full_composite_contract_matches_history_but_sibling_suffix_does_not():
    history = [service.UnifiedHistoryRow(row=4, contract_no='C1/EP27', reference_no='EDUHX00042', product_no='I1')]
    for contract, expected in [('C1/EP27', 'EDUHX00042'), ('C1/EP28', 'EDUHX00099')]:
        row = dict(contract_no=contract, po_no='', reference_no='EDUHX00099', product_no='I1', issues=[])
        mapping.prepare_row(row, history, 'edu', ['ITEM表'])
        assert row['reference_no'] == expected


def test_missing_customer_identity_is_not_replaced_by_allocated_edu_number():
    preview = service.create_unified_customer_preview(**preview_args(po_bytes([''])))
    assert preview['summary']['blocked'] == 1
    assert any(i['code'] == 'edu_missing_business_identity' for i in preview['rows'][0]['issues'])


def test_manual_contract_change_rechecks_business_identity_not_only_internal_reference():
    history = [service.UnifiedHistoryRow(row=4, contract_no='C1/EP1', reference_no='EDUHX00042', product_no='I1')]
    row = dict(id='r', contract_no='C1/EP1', po_no='', reference_no='EDUHX00099', product_no='I1', _huaxing_customer='edu')
    original = {'r': ('I1', 'EDUHX00099', 'C2/EP2', '')}
    with pytest.raises(service.CustomerOrderUnifiedError, match='重复'):
        mapping.validate_edited_identities({'rows': [row]}, original, history, ['ITEM表'])


def test_factory_scope_still_rejects_edu_in_huakang_a():
    args = preview_args(po_bytes(['C1']))
    with pytest.raises(service.CustomerOrderUnifiedError):
        service.create_unified_customer_preview(**{**args, 'factory_id': 'huakang-a'})


def test_multi_order_workbook_revision_does_not_drop_unrelated_orders():
    args = preview_args(po_bytes(['C1', 'C2'], pos=['P1', 'P2']))
    args['po_files'].append(('EDU_订单_REV2.xls', po_bytes(['C1', 'C3'], quantities=[6400, 6360], pos=['P1', 'P3'])))
    rows = service.create_unified_customer_preview(**args)['rows']
    assert {r['contract_no'] for r in rows} == {'C1', 'C2', 'C3'}
    assert next(r for r in rows if r['contract_no'] == 'C1')['quantity'] == '6400'
    assert len({r['reference_no'] for r in rows}) == 3


def test_history_quantity_conflict_cannot_be_confirmed_as_duplicate():
    wb = template('edu')
    sheet = wb['ITEM表']
    sheet['D4'], sheet['E4'], sheet['H4'], sheet['K4'] = 'C1/EP1', 'EDUHX00042', 'F12-MS907-B3080000', 6300
    args = preview_args(po_bytes(['C1/EP1']), wb)
    row = service.create_unified_customer_preview(**args)['rows'][0]
    assert row['reference_no'] == 'EDUHX00042'
    issue = next(i for i in row['issues'] if i['code'] == 'existing_quantity_conflict')
    assert issue['can_skip'] is False
    with pytest.raises(service.CustomerOrderUnifiedError):
        service.export_unified_customer_schedule(**args)


def test_batch_internal_reference_collision_is_not_confirmable():
    rows = [dict(id=f'r{i}', reference_no='EDUHX00042', contract_no=f'C{i}', product_no='I1', po_no='',
                 quantity='12', issues=[], _huaxing_customer='edu') for i in range(2)]
    service._mark_batch_duplicates(rows)
    issue = next(i for i in rows[1]['issues'] if i['code'] == 'edu_reference_collision')
    assert issue['can_skip'] is False


def test_same_revision_conflicting_files_are_not_silently_replaced():
    args = preview_args(po_bytes(['C1'], quantities=[6360]))
    args['po_files'].append(('EDU副本.xls', po_bytes(['C1'], quantities=[6400])))
    with pytest.raises(service.CustomerOrderUnifiedError, match='同版本内容冲突'):
        service.create_unified_customer_preview(**args)


def test_shared_history_internal_reference_for_other_contract_stays_blocked():
    wb = template('edu')
    for i, contract in enumerate(['C1', 'C2'], 4):
        sheet = wb['ITEM表']
        sheet.cell(i, 4, contract)
        sheet.cell(i, 5, 'EDUHX00042')
        sheet.cell(i, 8, 'F12-MS907-B3080000')
        sheet.cell(i, 11, 6360)
    row = service.create_unified_customer_preview(**preview_args(po_bytes(['C1']), wb))['rows'][0]
    assert any(i['code'] == 'edu_reference_collision' and not i['can_skip'] for i in row['issues'])
