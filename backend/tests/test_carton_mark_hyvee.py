import io

import pytest
from fastapi import HTTPException
from openpyxl import Workbook
from pypdf import PdfReader

from app.services.carton_mark_zuru import extract_zuru,prepare_records
from app.services.carton_supplier_mark_pdf import prepare_pdf
from test_carton_mark_zuru import original


def hyvee(changes=None,duplicate=False):
    book=Workbook();sheet=book.active;sheet.title="外箱数据"
    headers=["Hy-Vee PO#:","Factory PO","Factory","Unit UPC","Case UPC (UCN):","Customer Item NO.",
        "Amount of Inner cases:","Inner case pack:","ITEM NO.","Vendor Code:","Product Description:","Order Qty.","Master case pack:"]
    sheet.append(headers)
    row=["00357431","4500215651","华康","193052099723","10193052099720","00527590","4","6","00123-S001-PKC","00123","FUGGLER BADDIES","792","24"]
    for column,value in (changes or {}).items():row[headers.index(column)]=value
    sheet.append(row)
    sheet.append(row if duplicate else ["00357431","4500215651","华康","193052099891","10193052099898","00537270","N/A","N/A","00999-S002","00999","CARE BEAR KEYRINGS","144","6"])
    # Old exemplar is deliberately outside the data table. It must not enroll
    # an extra product or replace the authoritative current table values.
    sheet["Q8"]="Vendor:";sheet["R8"]="ZURU TOYS"
    sheet["Q14"]="Vendor Code:";sheet["R14"]="9639SQ1"
    sheet["Q16"]="Product Description:";sheet["R16"]="OLD EXAMPLE SHP COCO LOVE"
    sheet["S22"]="MADE IN CHINA"
    stream=io.BytesIO();book.save(stream);book.close();return stream.getvalue()


def source(**kwargs):
    result=original(excel_bytes=hyvee(),item_no="00123-S001-PKC")
    result["layout"]["config"]["paper"]="A3"
    result.update(kwargs);return result


def test_hyvee_data_rows_preserve_complete_sku_and_zeros_skip_stale_examples():
    records=extract_zuru("unseen.xlsx",hyvee())
    assert [r["fields"]["item"] for r in records]==["00123-S001-PKC","00999-S002"]
    assert records[0]["fields"]["customer_po"]=="00357431"
    assert records[0]["fields"]["article"]=="00527590"
    record,_=prepare_records(source(item_no="00999-S002"))
    assert record["fields"]["inner_pack"]=="N/A" and record["fields"]["quantity"]=="6"
    with pytest.raises(HTTPException):prepare_records(source(item_no="00123"))
    record,notes=prepare_records(source(item_no="00123",source_item="00123-S001-PKC"))
    assert record["fields"]["product"]=="FUGGLER BADDIES" and any("明确指定" in n for n in notes)


@pytest.mark.parametrize("column,value",[("Unit UPC","193052099724"),("Case UPC (UCN):","10193052099721"),
    ("Unit UPC","19305209972"),("Case UPC (UCN):","0193052099720"),("Unit UPC","#REF!"),
    ("Amount of Inner cases:","5"),("Inner case pack:","N/A"),("Master case pack:","0")])
def test_hyvee_invalid_barcode_checks_missing_fields_and_packaging_conflicts_reject(column,value):
    with pytest.raises(HTTPException) as error:extract_zuru("source.xlsx",hyvee({column:value}))
    assert error.value.status_code==422


def test_hyvee_exact_order_contract_and_duplicate_sku_reject():
    with pytest.raises(HTTPException):prepare_records(source(contract_no="wrong-order"))
    with pytest.raises(HTTPException):prepare_records(source(excel_bytes=hyvee(duplicate=True)))


@pytest.mark.parametrize("mode,paper,pages",[("front_side","A3",1),("front_only","A4",1),("separate_pages","A4",2)])
def test_hyvee_regenerates_two_barcode_numbers_fresh_text_and_respects_modes(mode,paper,pages):
    data=source();data["layout"]["config"].update(mode=mode,paper=paper)
    content,count,notes=prepare_pdf(data);reader=PdfReader(io.BytesIO(content))
    assert count==len(reader.pages)==pages
    text="\n".join(p.extract_text() for p in reader.pages)
    assert "10193052099720" in text and "193052099723" in text and "24 PC" in text
    assert "00527590" in text and "OLD EXAMPLE" not in text and "00999-S002" not in text
    assert any("内嵌旧条码" in n for n in notes)
    if mode!="front_only":assert "MADE IN CHINA" in text


def test_hyvee_barcode_layout_refuses_small_paired_page_instead_of_compressing():
    data=source();data["layout"]["config"]["paper"]="A4"
    with pytest.raises(HTTPException) as error:prepare_pdf(data)
    assert "A3" in error.value.detail


@pytest.mark.parametrize("case",["duplicate_header","country_conflict"])
def test_hyvee_duplicate_column_and_conflicting_origin_are_not_silently_chosen(case):
    from openpyxl import load_workbook
    book=load_workbook(io.BytesIO(hyvee()));sheet=book.active
    if case=="duplicate_header":sheet["N1"]="Unit UPC"
    else:sheet["S23"]="MADE IN VIETNAM"
    stream=io.BytesIO();book.save(stream);book.close()
    with pytest.raises(HTTPException):extract_zuru("source.xlsx",stream.getvalue())


def test_hyvee_api_saved_original_remains_order_bound_and_unapproved(monkeypatch):
    import hashlib,json,sqlalchemy as sa
    from test_molding_sample_api import make_client
    from test_carton_supplier_mark_pdf import prepare_layout,BASE
    from app.models.carton_mark import CartonMarkAsset,CartonMarkLayout,CartonMarkTemplate
    from app.models.carton_procurement import CartonAuditEvent
    from app.schemas.carton_supplier_portal import MarkLayoutConfig
    with make_client(monkeypatch) as client:
        order,payload,_=prepare_layout(client)
        from app.db import SessionLocal
        data=hyvee({"Factory PO":order["contract_no"]})
        with SessionLocal() as db:
            asset=db.get(CartonMarkAsset,payload["excel_asset_id"])
            asset.content=data;asset.size_bytes=len(data);asset.sha256=hashlib.sha256(data).hexdigest()
            layout=db.get(CartonMarkLayout,payload["layout_id"])
            layout.config_json=MarkLayoutConfig(series="zuru",paper="A3",logo_region=dict(x=10,y=10,width=20,height=20)).model_dump_json()
            db.commit()
        response=client.post(BASE,json={**payload,"source_item":"00123-S001-PKC"})
        assert response.status_code==201,response.text
        result=response.json();assert any("内嵌旧条码" in n for n in result["warnings"])
        assert result["asset"]["orders"][0]["id"]==order["id"]
        with SessionLocal() as db:
            saved=db.get(CartonMarkAsset,result["asset"]["id"])
            assert "10193052099720" in PdfReader(io.BytesIO(bytes(saved.content))).pages[0].extract_text()
            assert db.scalar(sa.select(sa.func.count()).select_from(CartonMarkTemplate))==0
            event=db.scalar(sa.select(CartonAuditEvent).where(CartonAuditEvent.entity_id==saved.id,CartonAuditEvent.event_type=="CARTON_MARK_ASSET_CREATED"))
            assert json.loads(event.detail_json)["conversion"]["source_item"]=="00123-S001-PKC"
