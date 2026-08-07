import hashlib
import importlib
import json
import os
import re
import sys
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select

TEST_TMP_DIR = Path(__file__).resolve().parents[1] / ".pytest-tmp"
BACKEND_DIR = Path(__file__).resolve().parents[1]
ADMIN_TEST_PASSWORD = "AdminSeed123!"

if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))


def make_client(monkeypatch) -> TestClient:
    TEST_TMP_DIR.mkdir(exist_ok=True)
    monkeypatch.setenv(
        "DATABASE_URL",
        f"sqlite:///{TEST_TMP_DIR / f'injection_phase4_{uuid4().hex}.db'}",
    )
    monkeypatch.setenv("SEED_ADMIN_PASSWORD", ADMIN_TEST_PASSWORD)
    monkeypatch.setenv("AUTHZ_MODE", "enforce")
    monkeypatch.setenv("AUTHZ_WRITES_ENABLED", "false")
    for module_name in list(sys.modules):
        if module_name == "app" or module_name.startswith("app."):
            del sys.modules[module_name]
    main = importlib.import_module("app.main")
    return TestClient(main.app)


def login(client: TestClient, username: str, password: str = "123456") -> dict:
    response = client.post(
        "/api/auth/login",
        json={"username": username, "password": password},
    )
    assert response.status_code == 200, response.text
    return response.json()


def ensure_user(
    username: str,
    role_id: str,
    *,
    factory_id: str,
    department: str,
) -> None:
    db_module = importlib.import_module("app.db")
    auth_models = importlib.import_module("app.models.auth")
    auth_service = importlib.import_module("app.services.auth")
    user_id = f"user-{username}"
    with db_module.SessionLocal() as db:
        salt, password_hash = auth_service.make_password_hash("123456")
        db.add(
            auth_models.AuthUser(
                id=user_id,
                username=username,
                display_name=username,
                password_salt=salt,
                password_hash=password_hash,
                status="active",
                force_password_change=0,
                created_at=auth_service.now_text(),
                updated_at=auth_service.now_text(),
            )
        )
        db.add(
            auth_models.AuthUserRole(
                id=f"{user_id}:{role_id}:{factory_id}:{department}",
                user_id=user_id,
                role_id=role_id,
                factory_id=factory_id,
                department=department,
            )
        )
        db.commit()


def build_workbook(*, missing_formula_cache: bool = False) -> bytes:
    workbook = Workbook()
    plan = workbook.active
    plan.title = "计划表"
    plan.append(["河源华兴啤机生产日计划表"])
    plan.append([])
    for column, value in {
        "B": "机号",
        "G": "工模",
        "H": "名称",
        "I": "单号",
        "J": "货号",
        "L": "订单数",
        "M": "已啤数",
        "O": "计划目标",
        "AG": "计划生产期",
        "AH": "计划完成期",
        "AR": "仓库",
    }.items():
        plan[f"{column}3"] = value
    plan["A4"] = "旧2"
    plan["B4"] = "旧2"
    plan["D4"] = "全自动"
    plan["G4"] = "32A 320T"
    plan["H4"] = "高速"
    plan["I4"] = "双臂五轴"
    plan["B5"] = "旧2"
    plan["D5"] = "全自动"
    plan["E5"] = "▲"
    plan["F5"] = "32A"
    plan["G5"] = "GT1214"
    plan["H5"] = "自卸车轮"
    plan["I5"] = 123
    plan["I5"].number_format = "000000"
    plan["J5"] = 45
    plan["J5"].number_format = "0000"
    plan["L5"] = 100
    plan["M5"] = "=5+5" if missing_formula_cache else 10
    plan["O5"] = 40
    plan["Q5"] = "黑色"
    plan["S5"] = "PP 5100NA"
    plan["T5"] = 379
    plan["Z5"] = date(2026, 7, 20)
    plan["AA5"] = date(2026, 8, 1)
    plan["AB5"] = date(2026, 8, 3)
    plan["AG5"] = datetime(2026, 8, 1, 8, 0)  # noqa: DTZ001
    plan["AH5"] = datetime(2026, 8, 1, 16, 0)  # noqa: DTZ001
    plan["AR5"] = 7
    plan["AR5"].number_format = "000"
    plan["AU5"] = "双臂"
    plan["AV5"] = "吸盘"

    mold = workbook.create_sheet("机安")
    for column, value in {
        "A": "模号",
        "B": "名称",
        "G": "安数",
        "H": "单双臂",
        "I": "吸盘",
        "L": "用料",
        "M": "工程重量（G）",
        "T": "L (mm)长/厚",
        "U": "W (mm)宽",
        "V": "H (mm)高",
    }.items():
        mold[f"{column}1"] = value
    mold["A2"] = "GT1214"
    mold["B2"] = "自卸车轮"
    mold["G2"] = "32A"
    mold["H2"] = "双臂"
    mold["I2"] = "吸盘"
    mold["L2"] = "PP"
    mold["M2"] = 379
    mold["T2"] = 550
    mold["U2"] = 450
    mold["V2"] = 850

    machine = workbook.create_sheet("华兴机器设备")
    machine["A2"] = "摆放区域"
    machine["B2"] = "机位"
    machine["A4"] = "老车间"
    machine["B4"] = 2
    machine["E4"] = "32A"
    machine["H4"] = 617
    machine["I4"] = "680MM*680MM"
    machine["J4"] = 320
    machine["L4"] = "高速"
    machine["M4"] = "五轴双臂"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_formula_error_workbook(error_value: str) -> bytes:
    source = build_workbook(missing_formula_cache=True)
    output = BytesIO()
    with ZipFile(BytesIO(source)) as original, ZipFile(
        output, "w", ZIP_DEFLATED
    ) as rewritten:
        for item in original.infolist():
            content = original.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                replacement = (
                    f'<c r="M5" t="e"><f>NA()</f><v>{error_value}</v></c>'
                ).encode()
                content, count = re.subn(
                    rb'<c r="M5"[^>]*>.*?</c>',
                    replacement,
                    content,
                    count=1,
                )
                assert count == 1
            rewritten.writestr(item, content)
    return output.getvalue()


def test_phase4_parser_preserves_identifiers_and_reports_missing_formula_cache():
    parser = importlib.import_module("app.services.injection_scheduling_excel")
    clean = build_workbook()
    normalized, issues = parser.parse_injection_scheduling_workbook(clean, "计划.xlsx")
    assert normalized["summary"]["task_count"] == 1
    assert normalized["tasks"][0]["order_no"] == "000123"
    assert normalized["tasks"][0]["item_no"] == "0045"
    assert normalized["tasks"][0]["warehouse_text"] == "007"
    assert normalized["tasks"][0]["source"]["sheet_name"] == "计划表"
    assert normalized["tasks"][0]["source"]["source_row"] == 5
    assert not any(item["blocking"] for item in issues)

    formula_source = build_workbook(missing_formula_cache=True)
    formula_hash = hashlib.sha256(formula_source).hexdigest()
    normalized, issues = parser.parse_injection_scheduling_workbook(
        formula_source, "公式缺缓存.xlsx"
    )
    assert normalized["source_file_hash"] == formula_hash
    assert normalized["summary"]["task_count"] == 0
    assert any(
        item["code"] == "FORMULA_CACHE_MISSING"
        and item["field_name"] == "completed_quantity"
        and item["blocking"]
        for item in issues
    )
    assert hashlib.sha256(formula_source).hexdigest() == formula_hash

    for error_value in ("#N/A", "#REF!"):
        error_source = build_formula_error_workbook(error_value)
        normalized, issues = parser.parse_injection_scheduling_workbook(
            error_source, f"公式错误-{error_value}.xlsx"
        )
        assert normalized["summary"]["task_count"] == 0
        assert any(
            item["code"] == "FORMULA_ERROR"
            and item["field_name"] == "completed_quantity"
            and item["raw_value"] == error_value
            and item["blocking"]
            for item in issues
        )


def test_phase4_import_batch_recovers_and_reidentifies_from_scoped_artifact(
    monkeypatch,
):
    source = build_workbook()
    source_hash = hashlib.sha256(source).hexdigest()
    with make_client(monkeypatch) as client:
        login(client, "admin", ADMIN_TEST_PASSWORD)
        preview = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "expected_revision": "0"},
            files={
                "file": (
                    "recoverable-plan.xlsx",
                    source,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "phase4-recoverable-preview-0001"},
        )
        assert preview.status_code == 201, preview.text
        batch = preview.json()
        assert batch["artifact_available"] is True
        assert batch["artifact_expires_at"]
        assert batch["preview_generation"] == 1
        assert "storage_key" not in batch

        recovered = client.get(
            f"/api/injection-scheduling/imports/{batch['id']}",
            params={"factory_id": "huaxing"},
        )
        assert recovered.status_code == 200, recovered.text
        assert recovered.json()["normalized_sha256"] == batch["normalized_sha256"]
        assert recovered.json()["artifact_available"] is True
        wrong_scope = client.get(
            f"/api/injection-scheduling/imports/{batch['id']}",
            params={"factory_id": "huakang-b"},
        )
        assert wrong_scope.status_code == 404

        retried = client.post(
            f"/api/injection-scheduling/imports/{batch['id']}/retry",
            json={
                "factory_id": "huaxing",
                "expected_revision": batch["revision"],
                "request_id": "phase4-reidentify-0001",
            },
        )
        assert retried.status_code == 200, retried.text
        retried_batch = retried.json()
        assert retried_batch["revision"] == batch["revision"] + 1
        assert retried_batch["preview_generation"] == 2
        assert retried_batch["source_file_hash"] == source_hash
        assert retried_batch["artifact_available"] is True

        replay = client.post(
            f"/api/injection-scheduling/imports/{batch['id']}/retry",
            json={
                "factory_id": "huaxing",
                "expected_revision": batch["revision"],
                "request_id": "phase4-reidentify-0001",
            },
        )
        assert replay.status_code == 200, replay.text
        assert replay.json()["idempotent_replay"] is True

        db_module = importlib.import_module("app.db")
        import_models = importlib.import_module(
            "app.models.injection_scheduling_import"
        )
        with db_module.SessionLocal() as db:
            artifact = db.scalar(
                select(import_models.InjectionSchedulingUploadArtifact).where(
                    import_models.InjectionSchedulingUploadArtifact.batch_id
                    == batch["id"]
                )
            )
        assert artifact is not None
        assert artifact.source_sha256 == source_hash
        assert bytes(artifact.payload_blob) == source


def test_phase4_preview_confirm_idempotency_lineage_and_factory_scope(monkeypatch):
    source = build_workbook()
    source_hash = hashlib.sha256(source).hexdigest()
    with make_client(monkeypatch) as client:
        login(client, "admin", ADMIN_TEST_PASSWORD)
        ensure_user(
            "phase4-clerk",
            "molding_clerk",
            factory_id="huaxing",
            department="molding",
        )
        client.post("/api/auth/logout")
        current = login(client, "phase4-clerk")
        assert "injection_scheduling:import" in current["permissions"]

        denied = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huakang-b", "expected_revision": "0"},
            files={"file": ("计划.xlsx", source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers={"x-request-id": "phase4-denied-preview"},
        )
        assert denied.status_code == 403

        preview = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "expected_revision": "0"},
            files={"file": ("计划.xlsx", source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers={"x-request-id": "phase4-preview-0001"},
        )
        assert preview.status_code == 201, preview.text
        batch = preview.json()
        assert batch["source_file_hash"] == source_hash
        assert batch["summary"]["task_count"] == 1
        assert batch["tasks"][0]["order_no"] == "000123"
        assert batch["tasks"][0]["item_no"] == "0045"
        assert batch["status"] == "PREVIEW"
        assert batch["revision"] == 1

        replay_preview = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "expected_revision": "0"},
            files={"file": ("计划.xlsx", source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers={"x-request-id": "phase4-preview-0001"},
        )
        assert replay_preview.status_code == 201
        assert replay_preview.json()["id"] == batch["id"]
        assert replay_preview.json()["idempotent_replay"] is True

        confirm_payload = {
            "factory_id": "huaxing",
            "expected_revision": 1,
            "expected_plan_revision": 0,
            "request_id": "phase4-confirm-0001",
            "confirm_mode": "create_draft",
            "business_date": "2026-08-01",
            "acknowledged_blocking_issue_ids": [],
        }
        stale_confirm = client.post(
            f"/api/injection-scheduling/imports/{batch['id']}/confirm",
            json={**confirm_payload, "expected_revision": 2},
        )
        assert stale_confirm.status_code == 409
        assert stale_confirm.json()["detail"]["current_revision"] == 1
        blocked_for_master = client.post(
            f"/api/injection-scheduling/imports/{batch['id']}/confirm",
            json=confirm_payload,
        )
        assert blocked_for_master.status_code == 409
        assert blocked_for_master.json()["detail"]["code"] == "CANONICAL_BATCH_NOT_READY"

        client.post("/api/auth/logout")
        login(client, "admin", ADMIN_TEST_PASSWORD)
        difference_keys = [
            f"{item['entity_type']}:{item['business_key']}"
            for item in batch["master_differences"]
        ]
        approved = client.post(
            f"/api/injection-scheduling/imports/{batch['id']}/master-differences/approve",
            json={
                "factory_id": "huaxing",
                "expected_revision": 1,
                "request_id": "phase4-master-approve-0001",
                "reason": "管理员核对来源机台与模具",
                "differences": difference_keys,
            },
        )
        assert approved.status_code == 200, approved.text
        assert approved.json()["revision"] == 2

        client.post("/api/auth/logout")
        login(client, "phase4-clerk")
        confirm_payload["expected_revision"] = 2
        confirm_payload["expected_action_fingerprint"] = approved.json()[
            "action_fingerprint"
        ]
        confirmed = client.post(
            f"/api/injection-scheduling/imports/{batch['id']}/confirm",
            json=confirm_payload,
        )
        assert confirmed.status_code == 200, confirmed.text
        result = confirmed.json()
        assert result["status"] == "CONFIRMED"
        assert result["revision"] == 3
        assert result["result"] == {
            "action_counts": {
                "CREATE_BASELINE_TASK": 1,
                "CREATE_ORDER": 1,
            },
            "action_fingerprint": approved.json()["action_fingerprint"],
            "backlog_without_tasks": True,
            "created_machines": 0,
            "created_molds": 0,
            "locked_baseline": True,
            "plan_created": True,
            "successor_created": False,
        }
        plan = client.get(
            "/api/injection-scheduling/plans/current",
            params={"factory_id": "huaxing"},
        )
        assert plan.status_code == 200, plan.text
        task = plan.json()["plan"]["tasks"][0]
        order = plan.json()["plan"]["orders"][0]
        assert order["order_no"] == "000123"
        assert order["item_no"] == "0045"
        assert task["import_batch_id"] == batch["id"]
        assert task["source_sheet_name"] == "计划表"
        assert task["source_row"] == 5
        assert task["source_file_hash"] == source_hash

        backlog = client.get(
            "/api/injection-scheduling/backlog",
            params={"factory_id": "huaxing"},
        )
        assert backlog.status_code == 200, backlog.text
        assert backlog.json()["items"] == []

        confirm_replay = client.post(
            f"/api/injection-scheduling/imports/{batch['id']}/confirm",
            json=confirm_payload,
        )
        assert confirm_replay.status_code == 200
        assert confirm_replay.json()["idempotent_replay"] is True

        merge_preview = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "expected_revision": "0"},
            files={"file": ("计划.xlsx", source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers={"x-request-id": "phase4-preview-merge"},
        )
        assert merge_preview.status_code == 201, merge_preview.text
        merge_batch = merge_preview.json()
        merged = client.post(
            f"/api/injection-scheduling/imports/{merge_batch['id']}/confirm",
            json={
                **confirm_payload,
                "expected_revision": 1,
                "expected_plan_revision": result["confirmed_plan_revision"],
                "expected_action_fingerprint": merge_batch["action_fingerprint"],
                "request_id": "phase4-confirm-merge",
                "confirm_mode": "merge_draft",
            },
        )
        assert merged.status_code == 200, merged.text
        assert merged.json()["confirmed_plan_revision"] == result["confirmed_plan_revision"]
        assert merged.json()["result"]["action_counts"] == {"SKIP_IDENTICAL": 1}

        client.post("/api/auth/logout")
        login(client, "admin", ADMIN_TEST_PASSWORD)
        published = client.post(
            f"/api/injection-scheduling/plans/{result['confirmed_plan_id']}/publish",
            json={
                "factory_id": "huaxing",
                "expected_revision": result["confirmed_plan_revision"],
                "request_id": "phase4-publish-0001",
            },
        )
        assert published.status_code == 200, published.text
        assert published.json()["plan"]["status"] == "PUBLISHED"

        next_preview = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "expected_revision": "0"},
            files={"file": ("计划.xlsx", source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers={"x-request-id": "phase4-preview-0002"},
        )
        assert next_preview.status_code == 201, next_preview.text
        next_batch = next_preview.json()
        next_confirm = client.post(
            f"/api/injection-scheduling/imports/{next_batch['id']}/confirm",
            json={
                **confirm_payload,
                "expected_revision": 1,
                "expected_action_fingerprint": next_batch["action_fingerprint"],
                "request_id": "phase4-confirm-0002",
            },
        )
        assert next_confirm.status_code == 200, next_confirm.text
        assert next_confirm.json()["confirmed_plan_id"] != result["confirmed_plan_id"]
        db_module = importlib.import_module("app.db")
        execution_models = importlib.import_module(
            "app.models.injection_scheduling_execution"
        )
        with db_module.SessionLocal() as db:
            original_published = db.scalar(
                select(execution_models.InjectionSchedulingPlan).where(
                    execution_models.InjectionSchedulingPlan.id
                    == result["confirmed_plan_id"]
                )
            )
            assert original_published.status == "PUBLISHED"
        assert hashlib.sha256(source).hexdigest() == source_hash


def test_canonical_formula_integrity_error_cannot_be_overridden(monkeypatch):
    source = build_workbook(missing_formula_cache=True)
    with make_client(monkeypatch) as client:
        login(client, "admin", ADMIN_TEST_PASSWORD)
        preview = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "expected_revision": "0"},
            files={"file": ("公式缺缓存.xlsx", source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers={"x-request-id": "phase4-preview-blocking"},
        )
        assert preview.status_code == 201, preview.text
        batch = preview.json()
        blocking_ids = [item["id"] for item in batch["issues"] if item["blocking"]]
        assert blocking_ids
        unacknowledged = client.post(
            f"/api/injection-scheduling/imports/{batch['id']}/confirm",
            json={
                "factory_id": "huaxing",
                "expected_revision": 1,
                "expected_plan_revision": 0,
                "request_id": "phase4-confirm-blocked",
                "confirm_mode": "create_draft",
                "business_date": "2026-08-01",
                "acknowledged_blocking_issue_ids": [],
            },
        )
        assert unacknowledged.status_code == 409
        assert (
            unacknowledged.json()["detail"]["code"]
            == "NON_OVERRIDABLE_CANONICAL_ERRORS"
        )
        assert unacknowledged.json()["detail"]["issue_ids"]


def test_canonical_confirm_requires_master_review_before_takeover(monkeypatch):
    source = build_workbook()
    with make_client(monkeypatch) as client:
        login(client, "admin", ADMIN_TEST_PASSWORD)
        preview = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "expected_revision": "0"},
            files={
                "file": (
                    "公共Profile预览.xlsx",
                    source,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "profile-preview-before-phase2-confirm"},
        )
        assert preview.status_code == 201, preview.text
        batch_id = preview.json()["id"]

        db_module = importlib.import_module("app.db")
        import_models = importlib.import_module(
            "app.models.injection_scheduling_import"
        )
        import_service = importlib.import_module(
            "app.services.injection_scheduling_import"
        )
        with db_module.SessionLocal() as db:
            record = db.get(import_models.InjectionSchedulingImportBatch, batch_id)
            normalized = json.loads(record.normalized_json)
            normalized["profile"]["profile_family"] = "huakang_b_daily_plan"
            unsigned = dict(normalized)
            unsigned.pop("normalized_sha256", None)
            normalized["normalized_sha256"] = import_service._payload_hash(unsigned)
            record.normalized_json = import_service._json(normalized)
            record.normalized_sha256 = normalized["normalized_sha256"]
            db.commit()

        confirm = client.post(
            f"/api/injection-scheduling/imports/{batch_id}/confirm",
            json={
                "factory_id": "huaxing",
                "expected_revision": 1,
                "expected_plan_revision": 0,
                "request_id": "profile-confirm-requires-phase2",
                "confirm_mode": "create_draft",
                "business_date": "2026-08-01",
                "acknowledged_blocking_issue_ids": [],
            },
        )
        assert confirm.status_code == 409
        assert (
            confirm.json()["detail"]["code"]
            == "CANONICAL_BATCH_NOT_READY"
        )


def test_canonical_confirm_rejects_profile_revision_that_is_no_longer_active(
    monkeypatch,
):
    source = build_workbook()
    with make_client(monkeypatch) as client:
        login(client, "admin", ADMIN_TEST_PASSWORD)
        preview = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "expected_revision": "0"},
            files={
                "file": (
                    "Profile版本变化.xlsx",
                    source,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "profile-preview-before-retire"},
        )
        assert preview.status_code == 201, preview.text
        batch_id = preview.json()["id"]

        db_module = importlib.import_module("app.db")
        import_models = importlib.import_module(
            "app.models.injection_scheduling_import"
        )
        with db_module.SessionLocal() as db:
            profile = db.get(
                import_models.InjectionSchedulingImportProfile,
                "isprofile-huaxing-daily-v1",
            )
            profile.status = "RETIRED"
            profile.lifecycle_revision += 1
            db.commit()

        confirm = client.post(
            f"/api/injection-scheduling/imports/{batch_id}/confirm",
            json={
                "factory_id": "huaxing",
                "expected_revision": 1,
                "expected_plan_revision": 0,
                "request_id": "profile-confirm-after-retire",
                "confirm_mode": "create_draft",
                "business_date": "2026-08-01",
                "acknowledged_blocking_issue_ids": [],
            },
        )
        assert confirm.status_code == 409
        assert confirm.json()["detail"]["code"] == "IMPORT_PROFILE_STALE"


@pytest.mark.skipif(
    not os.getenv("INJECTION_SCHEDULING_REAL_WORKBOOK"),
    reason="set INJECTION_SCHEDULING_REAL_WORKBOOK for local real-template regression",
)
def test_phase4_real_huaxing_workbook_read_only_regression(monkeypatch):
    parser = importlib.import_module("app.services.injection_scheduling_excel")
    source_path = Path(os.environ["INJECTION_SCHEDULING_REAL_WORKBOOK"])
    before_hash = hashlib.sha256(source_path.read_bytes()).hexdigest()
    before_stat = source_path.stat()
    normalized, issues = parser.parse_injection_scheduling_workbook(
        source_path.read_bytes(), source_path.name
    )
    after_stat = source_path.stat()
    assert normalized["source_file_hash"] == before_hash
    assert {"计划表", "机安"} <= set(normalized["sheet_names"])
    assert normalized["summary"]["profile_code"] == "huaxing_daily_plan_v1"
    assert normalized["summary"]["machine_count"] == 0
    assert normalized["summary"]["mold_count"] >= 3300
    assert normalized["summary"]["order_count"] >= 30
    assert normalized["summary"]["invalid_row_count"] >= 200
    assert any(item["code"] == "FORMULA_CACHE_MISSING" for item in issues)
    with make_client(monkeypatch) as client:
        login(client, "admin", ADMIN_TEST_PASSWORD)
        preview = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "expected_revision": "0"},
            files={
                "file": (
                    source_path.name,
                    source_path.read_bytes(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "phase4-real-workbook-preview"},
        )
        assert preview.status_code == 201, preview.text
        payload = preview.json()
        assert payload["source_file_hash"] == before_hash
        assert all(
            payload["summary"].get(key) == value
            for key, value in normalized["summary"].items()
        )
    assert hashlib.sha256(source_path.read_bytes()).hexdigest() == before_hash
    assert after_stat.st_size == before_stat.st_size
    assert after_stat.st_mtime_ns == before_stat.st_mtime_ns


@pytest.mark.skipif(
    not os.getenv("INJECTION_SCHEDULING_HUAKANG_B_WORKBOOK"),
    reason=(
        "set INJECTION_SCHEDULING_HUAKANG_B_WORKBOOK for local Huakang B "
        "real-template regression"
    ),
)
def test_phase4_real_huakang_b_workbook_read_only_preview(monkeypatch):
    source_path = Path(os.environ["INJECTION_SCHEDULING_HUAKANG_B_WORKBOOK"])
    source = source_path.read_bytes()
    before_hash = hashlib.sha256(source).hexdigest()
    before_stat = source_path.stat()
    with make_client(monkeypatch) as client:
        login(client, "admin", ADMIN_TEST_PASSWORD)
        preview = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huakang-b", "expected_revision": "0"},
            files={
                "file": (
                    source_path.name,
                    source,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "phase4-real-huakang-b-preview"},
        )
        assert preview.status_code == 201, preview.text
        payload = preview.json()
    summary = payload["summary"]
    assert payload["source_file_hash"] == before_hash
    assert payload["profile"]["profile_code"] == "huakang_b_daily_plan_v1"
    assert payload["batch_state"] == "MASTER_REVIEW_REQUIRED"
    assert summary["scheduled_baseline_count"] == 466
    assert summary["order_count"] == 454
    assert summary["invalid_row_count"] == 27
    assert summary["ignored_row_count"] == 98
    assert summary["issue_count"] == 126
    assert summary["blocking_issue_count"] == 69
    assert summary["master_difference_count"] == 419
    assert len(payload["reconciliation_actions"]) == 920
    after_stat = source_path.stat()
    assert hashlib.sha256(source_path.read_bytes()).hexdigest() == before_hash
    assert after_stat.st_size == before_stat.st_size
    assert after_stat.st_mtime_ns == before_stat.st_mtime_ns
