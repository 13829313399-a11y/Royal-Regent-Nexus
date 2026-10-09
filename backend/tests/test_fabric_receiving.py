import importlib
import json
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from threading import Barrier
from types import SimpleNamespace
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from test_fabric_procurement import BASE, book, row, upload, stored
from test_fabric_procurement_tracking import save, undo_preview, undo
from test_molding_sample_api import login_as
from warehouse_location_fixtures import make_client, seed_fabric_locations


def request(line, **changes):
    body = {"factory_id": "huakang-c", "request_id": str(uuid4()), "expected_source_revision": line["revision"],
        "expected_receipt_count": line.get("receipt_count", 0), "receipt_date": "2026-09-16",
        "delivery_reference": "DN-001", "material_category": "FABRIC",
        "batches": [{"quantity": "20.1", "location": "A01", "dye_lot": "00001", "roll_no": "00002"},
                    {"quantity": "29.9", "location": "A02", "dye_lot": "00001", "roll_no": "00003"}],
        "confirmed": True, **changes}
    for part in body["batches"]:
        part.setdefault("location_id", "fabric-" + part["location"].strip())
    return body


def receive(client, line, body):
    return client.post(f"{BASE}/lines/{line['id']}/receive", json=body)


def supplement_book(pending=True, returned=False):
    wb = load_workbook(BytesIO(book([row()])))
    headers = ["供应商复期", "放产日期", "订单号", "供应商", "生产单号", "款号", "物料编号", "旧编码", "物料名称", "订货量", "单位", "单价", "合同号", "交货明细", "入库数量", "车间组别", "送货单", "补数单号"]
    for name, enabled in (("补数未回物料", pending), ("补数已回", returned)):
        if not enabled:
            continue
        ws = wb.create_sheet(name)
        ws.append(headers)
        ws.append(["2026/9/20", "2026/9/01", "CGDD001", "采购供应商", "P001", "款001", "000123", "00123", "白色布 58寸", 100, "码", 0, "SC001", "100" if returned else "20+30", 100 if returned else 50, "车间A", "DN001", "00007"])
    out = BytesIO(); wb.save(out); return out.getvalue()


def test_supplement_tracking_preserves_separate_identity_metadata_and_returned_change(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        content = supplement_book()
        preview = upload(client, content, scope="TRACKING").json()
        assert preview["counts"]["new"] == 2
        facts = next(item["facts"] for item in preview["items"] if item["sheet"] == "补数未回物料")
        assert facts["source_category"] == "SUPPLEMENT" and facts["release_date"] == "2026-09-01"
        assert facts["order_date"] is None and facts["style_no"] == "款001"
        assert facts["supplement_no"] == "00007" and facts["workshop_group"] == "车间A" and facts["unit_price"] == "0"
        save(client, content)
        original = next(item for item in stored(client)["items"] if item["facts"].get("source_category") == "SUPPLEMENT")
        assert upload(client, content, scope="TRACKING").json()["counts"]["unchanged"] == 2
        changed = supplement_book(False, True)
        preview = upload(client, changed, scope="TRACKING").json()
        assert preview["counts"]["updated"] == 1
        save(client, changed)
        updated = next(item for item in stored(client)["items"] if item["id"] == original["id"])
        assert updated["facts"]["status"] == "RETURNED" and updated["tracking"]["arrival_review"]
        # Namespace remains distinct even without supplement number or style.
        parser = importlib.import_module("app.services.fabric_procurement_parser")
        values = {"order_no": "PO", "supplier": "S", "production_no": "P", "material_code": "M", "material_name": "N", "unit": "码", "ordered_quantity": "10"}
        assert parser.normalize_row(values, "PENDING")[1] != parser.normalize_row({**values, "source_category": "SUPPLEMENT"}, "PENDING")[1]
        assert parser.normalize_row({**values, "source_line_id": "K1"}, "PENDING")[1] != parser.normalize_row({**values, "source_line_id": "K1", "source_category": "SUPPLEMENT"}, "PENDING")[1]


def test_receipt_posts_only_current_physical_splits_and_retry_does_not_duplicate(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        batch = save(client, book([row()]))
        line = stored(client)["items"][0]
        before_withdraw = undo_preview(client, batch).json()
        stale_import = upload(client, book([row(订单号="NEW")]), scope="TRACKING").json()
        assert client.get(BASE + "/stock", params={"factory_id": "huakang-c"}).json()["total"] == 0
        body = request(line)
        posted = receive(client, line, body)
        assert posted.status_code == 200, posted.text
        assert posted.json()["quantity"] == "50.0" and posted.json()["stock_posted"] is True
        assert receive(client, line, body).json() == posted.json()
        stock = client.get(BASE + "/stock", params={"factory_id": "huakang-c"}).json()
        assert stock["total"] == 2
        assert {item["quantity"] for item in stock["items"]} == {"20.1", "29.9"}
        assert all(item["quality_status"] == "PENDING_INSPECTION" for item in stock["items"])
        assert any(item["roll_no"] == "00002" for item in stock["items"])
        current = stored(client)["items"][0]
        assert current["warehouse_received_quantity"] == "50.0" and current["prior_received_quantity"] == "50"
        assert current["warehouse_outstanding_quantity"] == "0.0"
        assert stored(client, view="OUTSTANDING")["total"] == 0
        assert undo_preview(client, batch).status_code == 409
        assert undo(client, batch, before_withdraw).status_code == 409
        from test_fabric_procurement import commit
        assert commit(client, book([row(订单号="NEW")]), stale_import, scope="TRACKING").status_code == 409
        assert receive(client, line, {**body, "request_id": str(uuid4())}).status_code == 409
        assert receive(client, line, {**body, "note": "changed"}).status_code == 409
        history = client.get(BASE + "/receipts", params={"factory_id": "huakang-c"}).json()
        assert history["total"] == 1 and history["items"][0]["facts"]["reported_received_quantity"] == "50"


def test_partial_receipts_freeze_history_and_allow_same_note_on_another_source(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        save(client, book([row(采购明细ID="K1", 入库数量=0, 交货明细=""), row(订单号="OTHER", 生产单号="P002", 入库数量=0, 交货明细="")]))
        lines = stored(client)["items"]
        first = next(line for line in lines if line["facts"]["order_no"] == "CGDD001")
        other = next(line for line in lines if line["facts"]["order_no"] == "OTHER")
        body = request(first, batches=[{"quantity": "0.1", "location": "B1", "dye_lot": "001"}])
        assert receive(client, first, body).status_code == 200
        # A note spanning two materials is valid; uniqueness is source + note.
        assert receive(client, other, request(other, prior_received_quantity="0")).status_code == 200
        current = stored(client, search="CGDD001")["items"][0]
        assert current["warehouse_received_quantity"] == "0.1" and current["warehouse_outstanding_quantity"] == "99.9"
        assert stored(client, view="PARTIAL", search="CGDD001")["total"] == 1
        next_body = request(current, delivery_reference="DN002", batches=[{"quantity": "0.2", "location": "B1", "dye_lot": "001"}])
        assert receive(client, current, next_body).status_code == 200
        assert stored(client, search="CGDD001")["items"][0]["warehouse_received_quantity"] == "0.3"
        # A changed source cannot rewrite receipt facts, and changed units cannot accumulate.
        save(client, book([row(采购明细ID="K1", 基本单位="米", 单价=8)]))
        changed = stored(client, search="CGDD001")["items"][0]
        assert changed["receipt_reconciliation_required"] is True
        assert receive(client, changed, request(changed, prior_received_quantity="0", delivery_reference="DN003")).status_code == 409
        history = client.get(BASE + "/receipts", params={"factory_id": "huakang-c", "search": "CGDD001"}).json()
        assert all(item["facts"]["unit"] == "码" and item["facts"]["unit_price"] == "3.99" for item in history["items"])


def test_invalid_or_unconfirmed_receipts_are_atomic_and_overage_requires_reason(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        save(client, book([row()]))
        line = stored(client)["items"][0]
        invalid = [request(line, confirmed=False), request(line, prior_received_quantity=""), request(line, delivery_reference="  "),
            request(line, receipt_date="2099-01-01"), request(line, accounting_month="2026-13"), request(line, accounting_month="2026-08"),
            request(line, batches=[{"quantity": "0", "location": "A"}]), request(line, batches=[{"quantity": "0.0000001", "location": "A"}]),
            request(line, batches=[{"quantity": "10", "location": " "}]), request(line, batches=[{"quantity": "10", "location": "A", "dye_lot": " "}]),
            request(line, batches=[{"quantity": "51", "location": "A", "dye_lot": "001"}]),
            request(line, batches=[{"quantity": "10", "location": "A", "dye_lot": "001", "roll_no": "001"}, {"quantity": "10", "location": "A", "dye_lot": "001", "roll_no": "001"}])]
        for body in invalid:
            response = receive(client, line, body)
            assert response.status_code == 422, response.text
        assert client.get(BASE + "/receipts", params={"factory_id": "huakang-c"}).json()["total"] == 0
        assert client.get(BASE + "/stock", params={"factory_id": "huakang-c"}).json()["total"] == 0
        body = request(line, difference_reason="供应商多送，现场点清", batches=[{"quantity": "51", "location": "A", "dye_lot": "001"}])
        assert receive(client, line, body).status_code == 200
        current = stored(client)["items"][0]
        assert receive(client, current, request(current, prior_received_quantity="49", delivery_reference="DN2")).status_code == 409


def test_imported_owed_quantity_and_later_reference_updates_do_not_double_subtract(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        saved = save(client, book([row(采购明细ID="K1", 订单数量=665, 入库数量=628, 交货明细="")]))
        line = stored(client)["items"][0]
        assert line["warehouse_outstanding_quantity"] == "37" and line["prior_received_quantity"] == "628"
        assert client.get(BASE + "/stock", params={"factory_id": "huakang-c"}).json()["total"] == 0
        # A refreshed procurement report is not warehouse confirmation.
        save(client, book([row(采购明细ID="K1", 订单数量=665, 入库数量=665, 交货明细="")]))
        line = stored(client)["items"][0]
        assert line["warehouse_outstanding_quantity"] == "37" and line["tracking"]["arrival_review"]
        posted = receive(client, line, request(line, batches=[{"quantity": "10", "location": "A", "dye_lot": "0001"}]))
        assert posted.status_code == 200, posted.text
        assert posted.json()["accounting_month"] == "2026-09"
        assert posted.json()["chase_reference"]["import_id"] == saved["id"]
        assert posted.json()["chase_reference"]["evidence_id"]
        current = stored(client)["items"][0]
        assert current["warehouse_outstanding_quantity"] == "27" and current["warehouse_received_quantity"] == "10"
        save(client, book([row(采购明细ID="K1", 订单数量=665, 入库数量=638, 交货明细="")]))
        current = stored(client)["items"][0]
        assert current["warehouse_outstanding_quantity"] == "27"
        assert receive(client, current, request(current, delivery_reference="DN2", batches=[{"quantity": "27", "location": "A", "dye_lot": "0001"}])).status_code == 200
        assert stored(client, view="OUTSTANDING")["total"] == 0
        assert stored(client)["items"][0]["warehouse_received_quantity"] == "37"


@pytest.mark.parametrize("category", ["ACCESSORY", "THREAD"])
def test_unknown_quantity_requires_source_check_and_explicit_zero_can_receive_non_fabric(monkeypatch, category):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        save(client, book([row(采购明细ID="K1", 入库数量="#VALUE!", 交货明细="")]))
        line = stored(client)["items"][0]
        assert line["prior_received_quantity"] is None and line["receipt_quantity_review_required"]
        detail = client.get(BASE + "/lines/" + line["id"], params={"factory_id": "huakang-c"}).json()
        assert detail["can_receive"] is True
        assert receive(client, line, request(line, prior_received_quantity="0")).status_code == 409
        correction = save(client, book([row(采购明细ID="K1", 入库数量=0, 交货明细="")]))
        line = stored(client)["items"][0]
        assert line["warehouse_outstanding_quantity"] == "100"
        preview = undo_preview(client, correction).json()
        assert undo(client, correction, preview).status_code == 200
        line = stored(client)["items"][0]
        assert line["prior_received_quantity"] is None
        save(client, book([row(采购明细ID="K1", 入库数量=0, 交货明细="")]))
        line = stored(client)["items"][0]
        assert receive(client, line, request(line, material_category=category)).status_code == 422
        body = request(line, receipt_date="2026-08-31", material_category=category, batches=[{"quantity": "10", "location": "A"}])
        posted = receive(client, line, body)
        assert posted.status_code == 200, posted.text
        assert posted.json()["accounting_month"] == "2026-08" and posted.json()["prior_received_quantity"] == "0"
        assert posted.json()["difference_reason"] == ""


@pytest.mark.parametrize("change", [dict(入库数量=49, 交货明细=""), dict(订单数量=49), dict(基本单位="米"), dict(物料名称="蓝色布")])
def test_chase_reference_changes_require_review_before_first_receipt(monkeypatch, change):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        save(client, book([row(采购明细ID="K1")]))
        save(client, book([row(采购明细ID="K1", **change)]))
        line = stored(client)["items"][0]
        assert line["receipt_quantity_review_required"] or line["receipt_reconciliation_required"]
        assert line["warehouse_outstanding_quantity"] is None
        detail = client.get(BASE + "/lines/" + line["id"], params={"factory_id": "huakang-c"}).json()
        assert detail["can_receive"] is False
        assert receive(client, line, request(line)).status_code in (409, 422)
        assert client.get(BASE + "/stock", params={"factory_id": "huakang-c"}).json()["total"] == 0


def test_ambiguous_legacy_import_order_cannot_choose_a_random_starting_reference(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        save(client, book([row(采购明细ID="K1")]))
        save(client, book([row(采购明细ID="K1", 入库数量=60, 交货明细="")]))
        dbm = importlib.import_module("app.db")
        model = importlib.import_module("app.models.fabric_procurement")
        with dbm.SessionLocal() as db:
            for batch in db.scalars(select(model.FabricProcurementImport)):
                metadata = json.loads(batch.result_json)
                metadata.pop("sequence", None)
                batch.result_json = json.dumps(metadata)
                batch.occurred_at = "2026-09-16T08:00:00+08:00"
            db.commit()
        line = stored(client)["items"][0]
        assert line["receipt_quantity_review_required"] and line["warehouse_outstanding_quantity"] is None
        assert receive(client, line, request(line)).status_code == 422


@pytest.mark.parametrize("mode", ["legacy", "shadow", "enforce"])
def test_receiving_auth_factory_department_and_explicit_deny(monkeypatch, mode):
    monkeypatch.setenv("AUTHZ_MODE", mode)
    monkeypatch.setenv("AUTHZ_WRITES_ENABLED", "false")
    with make_client(monkeypatch) as client:
        assert importlib.import_module("app.core.config").settings.authz_mode == mode
        assert client.get(BASE + "/stock", params={"factory_id": "huakang-c"}).status_code == 401
        login_as(client, "admin")
        save(client, book([row()]))
        line = stored(client)["items"][0]
        assert receive(client, line, request(line, factory_id="huakang-d")).status_code == 422
        from test_raw_material_api import create_fixed_position_user, login_fixed_position_user
        create_fixed_position_user("receiver_c", "position_warehouse_keeper", "pmc-warehouse", factory_id="huakang-c")
        create_fixed_position_user("receiver_d", "position_warehouse_keeper", "pmc-warehouse", factory_id="huakang-d")
        login_fixed_position_user(client, "receiver_d")
        assert receive(client, line, request(line)).status_code == 403
        assert client.get(BASE + "/receipts", params={"factory_id": "huakang-c"}).status_code == 403
        login_fixed_position_user(client, "receiver_c")
        assert receive(client, line, request(line)).status_code == 200
        auth = importlib.import_module("app.services.auth")
        api = importlib.import_module("app.api.fabric_procurement")
        dbm = importlib.import_module("app.db")
        models = importlib.import_module("app.models.auth")
        from dataclasses import replace
        with dbm.SessionLocal() as db:
            actor = auth.build_auth_context(db, db.get(models.AuthUser, "user-receiver_c"))
            denied = replace(actor, overrides=(auth.AuthOverrideContext(id="deny-receive", permission_code="fabric_warehouse:receive", effect="deny", factory_id="huakang-c", department="pmc-warehouse"),))
            with pytest.raises(Exception) as refused:
                api.authorize(db, denied, "huakang-c", "receive")
            assert refused.value.status_code == 403
            wrong_department = replace(actor, grants=tuple(replace(grant, department="engineering", unrestricted_department=True) for grant in actor.grants))
            assert not auth.can(wrong_department, "fabric_warehouse:receive", "huakang-c", "pmc-warehouse")


def test_missing_receipt_schema_blocks_posting_without_disabling_sources(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        save(client, book([row()]))
        line = stored(client)["items"][0]
        dbm = importlib.import_module("app.db")
        receiving = importlib.import_module("app.services.fabric_receiving")
        for model in reversed(receiving.MODELS):
            model.__table__.drop(dbm.engine)
        assert stored(client)["total"] == 1
        response = receive(client, line, request(line))
        assert response.status_code == 503 and "20261008_0143" in response.text


@pytest.mark.parametrize("same_request", [True, False])
def test_concurrent_receiving_posts_once_and_rejects_stale_count(tmp_path, same_request):
    from fastapi import HTTPException
    service = importlib.import_module("app.services.fabric_receiving")
    source = importlib.import_module("app.services.fabric_procurement")
    parser = importlib.import_module("app.services.fabric_procurement_parser")
    schemas = importlib.import_module("app.schemas.fabric_receiving")
    models = importlib.import_module("app.models.fabric_procurement")
    dbm = importlib.import_module("app.db")
    engine = create_engine("sqlite:///" + str(tmp_path / "concurrent-receipts.db"), connect_args={"check_same_thread": False})
    tables = [m.__table__ for m in (models.FabricProcurementState, models.FabricProcurementLine, models.FabricProcurementImport, models.FabricProcurementEvidence, *service.MODELS)]
    dbm.Base.metadata.create_all(engine, tables=tables)
    actor = SimpleNamespace(id="actor", display_name="Warehouse")
    with Session(engine) as db:
        seed_fabric_locations(db)
        parsed = parser.parse_workbook("source.xlsx", book([row()]))
        token = source.preview(db, parsed, actor.id)["preview_token"]
        source.apply(db, parser.parse_workbook("source.xlsx", book([row()])), actor, "source.xlsx", str(uuid4()), token, confirmed=True, acknowledge_excluded=True)
        line = source.list_lines(db, "ALL")["items"][0]
    bodies = [request(line), request(line)]
    if same_request: bodies[1] = bodies[0]
    barrier = Barrier(2)
    def post(body):
        with Session(engine) as db:
            barrier.wait(timeout=10)
            try: return service.receive(db, actor, line["id"], schemas.ReceiveSourceRequest(**body))
            except HTTPException as exc: return exc.status_code
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(post, bodies))
    if same_request: assert results[0] == results[1] and isinstance(results[0], dict)
    else: assert sum(result == 409 for result in results) == 1 and sum(isinstance(result, dict) for result in results) == 1
    with Session(engine) as db:
        assert db.scalar(select(func.count()).select_from(service.FabricReceipt)) == 1
        assert db.scalar(select(func.count()).select_from(service.FabricInventoryMovement)) == 2


def test_receipt_migration_constraints_and_downgrade_guard(tmp_path):
    import sqlalchemy as sa
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from importlib.util import spec_from_file_location, module_from_spec
    def migration(filename):
        spec = spec_from_file_location("migration_" + filename, Path(__file__).parents[1] / "alembic/versions" / filename)
        module = module_from_spec(spec); spec.loader.exec_module(module); return module
    prior = migration("20261005_0134_fabric_procurement.py")
    current = migration("20261008_0142_fabric_receiving.py")
    engine = sa.create_engine("sqlite:///" + str(tmp_path / "receipt-migration.db"))
    with engine.begin() as connection:
        with Operations.context(MigrationContext.configure(connection)):
            prior.upgrade(); current.upgrade()
            assert set(("fabric_receipts", "fabric_stock_batches", "fabric_inventory_movements")) <= set(sa.inspect(connection).get_table_names())
            columns = {column["name"] for column in sa.inspect(connection).get_columns("fabric_receipts")}
            assert "source_revision" in columns and "prior_received_quantity" in columns
            connection.execute(sa.text("INSERT INTO fabric_procurement_lines VALUES ('L', 'huakang-c', 'KEY', 1, 'PENDING', 'PO', 'Supplier', 'MAT', 'PROD', '{}', 1, 'NOW')"))
            connection.execute(sa.text("INSERT INTO fabric_receipts VALUES ('R','huakang-c','L',1,'REQUEST','HASH','DN','2026-09-16','2026-09','1','码','0','{}','{}','A','Warehouse','NOW')"))
            with pytest.raises(RuntimeError, match="真实布料"):
                current.downgrade()
            connection.execute(sa.text("DELETE FROM fabric_receipts"))
            current.downgrade()
            assert "fabric_receipts" not in sa.inspect(connection).get_table_names()
