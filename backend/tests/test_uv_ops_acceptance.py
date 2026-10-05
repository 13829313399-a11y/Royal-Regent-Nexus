from dataclasses import replace
from decimal import Decimal
from io import BytesIO
from uuid import uuid4
import asyncio
from concurrent.futures import ThreadPoolExecutor
from openpyxl import load_workbook,Workbook
import pytest
from sqlalchemy.orm import Session
from sqlalchemy import select
from test_uv_ops import client,setup_task,post,read,payload,confirm,user_with
from test_uv_ops_edges import setup_agent,event
from app.models import uv_operations as m
from app.services.uv_operations import common as c,live,ingest,reports,exports,imports,runtime


def test_workspace_pg_json_null_and_sensitive_projection(client):
    data=setup_task(client)
    with Session(client.uv_engine) as db:
        task=db.get(m.UvOpsTask,data['task']['id']);task.cost_price_snapshot=None;db.commit()
    response=client.get('/api/uv-operations/workspace',params={'factory_id':'huakang-a'})
    assert response.status_code==200,response.text
    assert response.json()['meta']['coverage']['missing_cost_records']==1
    confirm(client,data,good=10,rework=0)
    report=client.get('/api/uv-operations/reports',params=dict(factory_id='huakang-a',start_date='2026-09-23',end_date='2026-09-23')).json()['data']
    assert all(value['cost_margin'] is None and value['cost_total'] is None for value in report['cost_by_currency'].values())
    client.uv_auth['user']=user_with('read','cost_read')
    report=client.get('/api/uv-operations/reports',params=dict(factory_id='huakang-a',start_date='2026-09-23',end_date='2026-09-23')).json()['data']
    assert all('cost_total' not in value and 'cost_margin' not in value and 'payroll_total' not in value for value in report['cost_by_currency'].values())


def test_agent_poison_events_old_stream_revocation_and_source_uniqueness(client,monkeypatch):
    machine,created,source,headers=setup_agent(client)
    bad=event(machine,source)|dict(sequence=2**50)
    poison=event(machine,source,2);poison['payload']['file_hash']='x'*300
    valid=event(machine,source,3)
    def upload(events):return client.post('/api/internal/uv-agent/events/batch',headers=headers,json=dict(schema_version=1,batch_id=uuid4().hex,events=events))
    assert [row['status'] for row in upload([bad,poison,valid]).json()['results']]==['rejected','rejected','persisted']
    switched=event(machine,source,4)|dict(stream_id='other-stream')
    assert upload([switched]).json()['results'][0]['code']=='stream_rebind_required'
    other=post(client,'machines',code='OTHER',name='合成第二机台')
    response=client.post('/api/uv-operations/sources',json=payload(agent_id=created['agent']['id'],machine_id=other['id'],source_id=source['source_id'],adapter_type='simulator'))
    assert response.status_code==409
    original=ingest.event
    def revoke_between(db,agent_id,raw):
        result=original(db,agent_id,raw)
        with Session(client.uv_engine) as separate:
            separate.get(m.UvOpsAgent,agent_id).revoked=True;separate.commit()
        return result
    monkeypatch.setattr(ingest,'event',revoke_between)
    rows=upload([event(machine,source,5),event(machine,source,6)]).json()['results']
    assert [row['status'] for row in rows]==['persisted','rejected']
    assert rows[1]['code']=='agent_revoked'
    assert client.post('/api/internal/uv-agent/heartbeat',headers=headers,json={}).status_code==401


def test_late_evidence_is_separate_and_closed_snapshot_is_immutable(client):
    machine,created,source,headers=setup_agent(client)
    with Session(client.uv_engine) as db:
        period=c.lock_period(db,'2026-09-01');identifier=period.id;db.commit()
    period=next(x for x in read(client,'periods') if x['id']==identifier)
    closed=post(client,'periods/2026-09/close',expected_version=period['version'],reason='合成空账期完整封存')
    body=dict(schema_version=1,batch_id=uuid4().hex,events=[event(machine,source)])
    ack=client.post('/api/internal/uv-agent/events/batch',headers=headers,json=body)
    assert ack.json()['results'][0]['disposition']=='late_evidence_review'
    assert read(client,'runs')==[]
    late=read(client,'late-events')[0]
    value=client.get('/api/uv-operations/reports',params=dict(factory_id='huakang-a',start_date='2026-09-01',end_date='2026-09-30')).json()['data']
    assert value['report_basis']=='closed_snapshot' and value['coverage']['unmatched_runs']==0
    post(client,'late-events/'+late['id']+'/acknowledge',expected_version=late['version'],reason='核对为补传设备证据，不变更生产数量')
    assert read(client,'late-events')[0]['resolved_by']


def test_future_pending_cannot_be_resolved_in_prior_business_day(client):
    data=setup_task(client,quantity=5)
    earlier=data['shift']
    data['shift']=post(client,'shifts',name='后日合成班',business_date='2026-09-24',start_at='2026-09-24T08:00:00+08:00',end_at='2026-09-24T16:00:00+08:00')
    result=confirm(client,data,good=0,rework=0,pending=5)
    response=client.post('/api/uv-operations/quality/resolve',json=payload(batch_id=data['batch']['id'],shift_id=earlier['id'],quantity=5,disposition='good',expected_version=result['batch']['version'],evidence='合成不合法提前处置'))
    assert response.status_code==409 and response.json()['code']=='future_source'


def test_rework_pending_quality_pays_actual_worker_and_not_parent(client):
    data=setup_task(client)
    root=data['task'];result=confirm(client,data);data['batch']=result['batch']
    child=post(client,'rework-tasks',batch_id=data['batch']['id'],expected_version=data['batch']['version'],code='CHILD',quantity=6)
    data['task']=child
    result=confirm(client,data,good=0,rework=0,scrap=1,pending=5,source_bucket='rework')
    quality=post(client,'quality/resolve',batch_id=data['batch']['id'],shift_id=data['shift']['id'],quantity=5,disposition='good',expected_version=result['batch']['version'],evidence='合成返工待检转良')
    assert quality['entry']['task_id']==child['id']
    for task,worker,expected in [(root,'worker-root',114),(child,'worker-child',5)]:
        post(client,'participations',shift_id=data['shift']['id'],task_id=task['id'],employee_id=worker,employee_name=worker,start_at='2026-09-23T08:00:00+08:00',end_at='2026-09-23T09:00:00+08:00')
        wage=post(client,'wages/confirm',task_id=task['id'],shift_id=data['shift']['id'])
        assert wage['accrual']['payroll_evidence']['final_good']==expected


def test_export_is_full_immutable_numeric_and_revocation_safe(client):
    data=setup_task(client)
    confirm(client,data,good=10,rework=0)
    job=post(client,'exports',kind='daily',start_date='2026-09-23',end_date='2026-09-23')
    with Session(client.uv_engine) as db:
        row=db.get(m.UvOpsExportJob,job['id']);exports.build(db,client.uv_auth['user'],row);db.commit()
        before=row.artifact
    response=client.get('/api/uv-operations/exports/'+job['id']+'/download',params=dict(factory_id='huakang-a'))
    assert response.status_code==200 and response.content==before
    book=load_workbook(BytesIO(response.content));sheet=book['UV记录']
    names=[cell.value for cell in sheet[1]]
    assert sheet.cell(2,names.index('currency')+1).value=='CNY'
    assert sheet.cell(2,names.index('cost_revenue')+1).data_type=='n'
    assert '按币种核算' in book.sheetnames and '数据缺项' in book.sheetnames
    client.uv_auth['user']=user_with('read','export')
    assert client.get('/api/uv-operations/exports/'+job['id']+'/download',params=dict(factory_id='huakang-a')).status_code==403


def test_import_previews_business_conflicts_without_writing_formal_rows(client):
    content=b'code,name\nDUP,first\nDUP,second\n'
    response=client.post('/api/uv-operations/imports/preview',data=dict(factory_id='huakang-a',template='products-v1',units_confirmed='true'),files={'file':('duplicate.csv',content,'text/csv')})
    assert response.status_code==200,response.text
    result=response.json()['data'];assert result['errors'][0]['row']==3
    assert read(client,'products')==[]
    assert client.post('/api/uv-operations/imports/'+result['id']+'/commit',json=payload(expected_version=result['version'])).status_code==409


def test_xlsx_declared_dimension_rejected_before_iteration():
    from zipfile import ZipFile,ZIP_DEFLATED
    book=Workbook();book.active.append(['code','name']);stream=BytesIO();book.save(stream)
    altered=BytesIO()
    with ZipFile(stream) as source,ZipFile(altered,'w',ZIP_DEFLATED) as target:
        for name in source.namelist():
            value=source.read(name)
            if name=='xl/worksheets/sheet1.xml':value=value.replace(b'A1:B1',b'A1:XFD1048576')
            target.writestr(name,value)
    with pytest.raises(c.DomainError) as error:
        imports.read_file('synthetic.xlsx',altered.getvalue())
    assert error.value.body['code']=='worksheet_dimensions'
