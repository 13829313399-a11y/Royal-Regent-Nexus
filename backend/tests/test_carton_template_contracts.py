"""Published empty templates must expose the current parser's business inputs."""
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from app.services.carton_master_import import HEADERS, parse, template
from app.services.carton_opening_preview import options_from_json, parse_rows as parse_opening
from app.services.carton_procurement_history_import import _parse_rows
from app.services.carton_procurement_imports import _parse_weekly
from fastapi import HTTPException

TEMPLATES = Path(__file__).resolve().parents[2] / "public/templates"


def saved(book):
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def test_history_template_preserves_positions_and_exposes_customer_po_without_guessing_product():
    book = load_workbook(TEMPLATES / "carton-history-order-import-template.xlsx")
    main = book["历史订单导入"]
    headers = [cell.value for cell in main[2]]
    assert headers[:8] == ["客户名称", "合同号", "货号", "下单日期", "计划交期", "外箱需求数量", "内箱需求数量", "卡纸需求数量"]
    assert headers[8:10] == ["外箱每箱净重kg", "外箱每箱毛重kg"]
    assert headers[-1] == "客户PO"
    assert not any(cell.value is not None for row in main.iter_rows(min_row=3) for cell in row)
    for target, po in [(3, "0001"), (4, "0002")]:
        for column, cell in enumerate(book["填写示例"][6], 1):
            main.cell(target, column, cell.value)
        main.cell(target, headers.index("客户PO") + 1, po)
    rows, _ = _parse_rows("history.xlsx", saved(book))
    assert [row["customer_po"] for row in rows] == ["0001", "0002"]
    assert {row["item_no"] for row in rows} == {"000203307004"}
    assert all(row["product_name"] == "" for row in rows)
    assert all(row["product_order_quantity"] is None for row in rows)
    assert all(row["line"].net_weight_kg == Decimal("8.125") for row in rows)
    assert book["字段说明"]["C16"].value == "保留原计划交期；缺客户交期时按计划交期加有效采购安全天数补齐，预览显示实际规则。"


def test_history_supplement_example_has_the_same_weight_headers_as_its_input_sheet():
    book = load_workbook(TEMPLATES / "carton-history-order-import-template.xlsx")
    expected = [cell.value for cell in book["附加纸品明细"][2]]
    actual = [cell.value for cell in book["填写示例"][14]][:len(expected)]
    assert actual == expected
    for column, cell in enumerate(book["填写示例"][8], 1):
        book["历史订单导入"].cell(3, column, cell.value)
    for target, source in [(3, 15), (4, 16)]:
        for column in range(1, len(expected) + 1):
            book["附加纸品明细"].cell(target, column, book["填写示例"].cell(source, column).value)
    rows, _ = _parse_rows("history.xlsx", saved(book))
    assert [row["line"].packaging_type for row in rows] == ["外箱", "平卡", "刀卡"]
    assert all(row["product_name"] == "" for row in rows)
    assert all(row["line"].net_weight_kg is None for row in rows[1:])


@pytest.mark.parametrize("kind", list(HEADERS))
def test_master_templates_examples_parse_with_current_contracts_and_empty_data_stays_empty(kind):
    book = load_workbook(BytesIO(template(kind)))
    assert [cell.value for cell in book["导入数据"][1]] == HEADERS[kind]
    assert book["导入数据"].max_row == 1
    examples = list(book["说明与示例"].iter_rows(values_only=True))
    header_index = next(i for i, row in enumerate(examples) if list(row) == HEADERS[kind])
    for row in examples[header_index + 1:]:
        book["导入数据"].append(row)
    entries, errors = parse(saved(book), "master.xlsx", kind)
    assert entries and errors == []
    if kind == "configurations":
        assert {entry[1]["code"] for entry in entries} == {"00123"}
        assert all(entry[1]["line"]["net_weight_kg"] is None for entry in entries)


def test_opening_template_examples_parse_without_turning_missing_price_into_zero():
    book = load_workbook(TEMPLATES / "carton-history-inventory-import-template.xlsx")
    sheet = book["期初库存导入"]
    assert sheet["A2"].value == "合同号"
    assert sheet["P2"].value == "客户"
    assert book["填写示例"]["A5"].value == "合同号"
    assert book["填写示例"]["P5"].value == "客户"
    assert sheet["O2"].value == "入库时间"
    for target, source in enumerate(range(6, 10), 3):
        for column, cell in enumerate(book["填写示例"][source], 1):
            sheet.cell(target, column, cell.value)
    rows, _ = parse_opening("opening.xlsx", saved(book), options_from_json({"customer_name": "迪奇", "warehouse": "A", "dimension_unit": "cm"}))
    assert len(rows) == 4
    assert rows[1]["unit_price"] is None and rows[1]["amount"] is None
    assert rows[-1]["status"] == "ZERO"
    assert all(row["item_no"].startswith("000") for row in rows)


@pytest.mark.parametrize("default_customer", ["", "迪奇"])
def test_opening_template_customer_column_is_optional_and_row_value_overrides_default(default_customer):
    book = load_workbook(TEMPLATES / "carton-history-inventory-import-template.xlsx")
    sheet = book["期初库存导入"]
    for target, source in enumerate(range(6, 9), 3):
        for column, cell in enumerate(book["填写示例"][source], 1):
            sheet.cell(target, column, cell.value)
    sheet["P3"] = "360"
    sheet["P4"] = None
    sheet["P5"] = "迪奇"
    rows, _ = parse_opening("opening.xlsx", saved(book), options_from_json({
        "customer_name": default_customer, "warehouse": "A", "dimension_unit": "cm",
    }))
    assert [row["customer_name"] for row in rows] == ["360", default_customer, "迪奇"]
    assert all(row["contract_no"] == "0010001234/200" for row in rows)
    assert all(row["item_no"] == "000203307004" for row in rows)

    # A legacy file without the appended customer column and using PO still works.
    sheet.delete_cols(16)
    sheet["A2"] = "PO"
    legacy, _ = parse_opening("legacy.xlsx", saved(book), options_from_json({
        "customer_name": default_customer, "warehouse": "A", "dimension_unit": "cm",
    }))
    assert len(legacy) == 3 and all(row["customer_name"] == default_customer for row in legacy)
    assert [row["contract_no"] for row in legacy] == [row["contract_no"] for row in rows]


def test_weekly_template_keeps_customer_po_separate_from_po_source_display():
    book = load_workbook(TEMPLATES / "carton-weekly-schedule-template.xlsx")
    main = book["导入数据"]
    values = ["0010001234/200", "0001", "DISPLAY-ONLY", "迪奇", "000203307004", "", 3600, "1/120", "2026-09-21"]
    for column, value in enumerate(values, 1):
        main.cell(4, column, value)
    parsed = _parse_weekly("schedule.xlsx", saved(book))
    assert len(parsed["rows"]) == 1
    assert parsed["rows"][0]["customer_po"] == "0001"
    assert parsed["rows"][0]["po_numbers"] == "DISPLAY-ONLY"
    assert parsed["rows"][0]["item_no"] == "000203307004"


def test_published_empty_templates_without_ooxml_dimensions_report_no_data_instead_of_crashing():
    with pytest.raises(HTTPException, match="没有可导入"):
        _parse_rows("history.xlsx", (TEMPLATES / "carton-history-order-import-template.xlsx").read_bytes())
    with pytest.raises(HTTPException, match="未找到可导入数据"):
        parse_opening("opening.xlsx", (TEMPLATES / "carton-history-inventory-import-template.xlsx").read_bytes(), options_from_json({"dimension_unit": "cm"}))


@pytest.mark.parametrize("strict", [False, True])
def test_missing_dimensions_still_enforces_existing_physical_row_limits(monkeypatch, strict):
    from app.services import carton_procurement_imports as imports
    from openpyxl import Workbook
    book = Workbook()
    book.active.title = "ITEM" if strict else "导入数据"
    for number in range(5):
        book.active.append([str(number)])
    output = BytesIO()
    with ZipFile(BytesIO(saved(book))) as source, ZipFile(output, "w") as target:
        for member in source.infolist():
            data = source.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb"<dimension\b[^>]*/>", b"", data)
            target.writestr(member, data)
    monkeypatch.setattr(imports, "MAX_IMPORT_ROWS", 3)
    with pytest.raises(HTTPException, match="超过"):
        imports._sheet_rows("source.xlsx", output.getvalue(), strict_item_limit=strict)
