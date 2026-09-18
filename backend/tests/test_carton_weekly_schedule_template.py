from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from test_carton_procurement_api import _create_order, _freeze_carton_time
from test_molding_sample_api import login_as, make_client


TEMPLATE_PATH = (
    Path(__file__).resolve().parents[2]
    / "public"
    / "templates"
    / "carton-weekly-schedule-template.xlsx"
)


def test_weekly_schedule_template_matches_import_contract(monkeypatch):
    workbook = load_workbook(TEMPLATE_PATH)
    assert workbook.sheetnames == ["导入数据", "填写说明与示例"]

    data_sheet = workbook["导入数据"]
    assert [data_sheet.cell(row=3, column=index).value for index in range(1, 10)] == [
        "Reference",
        "客户采购单号",
        "PO.NO",
        "客户",
        "货号",
        "产品名称",
        "数量",
        "装箱",
        "验货期",
    ]
    assert data_sheet.max_row >= 203
    assert all(data_sheet.cell(row=4, column=index).value is None for index in range(1, 10))
    assert "系统怎么核对" in str(workbook["填写说明与示例"]["A14"].value)

    data_sheet.cell(row=4, column=1, value="SC700145365")
    data_sheet.cell(row=4, column=3, value="PO-001")
    data_sheet.cell(row=4, column=4, value="Dickie")
    data_sheet.cell(row=4, column=5, value="203302044")
    data_sheet.cell(row=4, column=6, value="多文盒")
    data_sheet.cell(row=4, column=7, value=3600)
    data_sheet.cell(row=4, column=8, value="1/120")
    data_sheet.cell(row=4, column=9, value="2026-09-21 至 2026-09-24")

    content = BytesIO()
    workbook.save(content)
    workbook.close()

    with make_client(monkeypatch) as client:
        login_as(client, "warehouse_keeper")
        _freeze_carton_time(monkeypatch)
        _create_order(client)
        response = client.post(
            "/api/carton-procurement/weekly-imports",
            params={"factory_id": "huaxing"},
            files={
                "file": (
                    "纸箱每周排期核对导入模板.xlsx",
                    content.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

        assert response.status_code == 201, response.text
        summary = response.json()["parse_summary"]
        assert summary["row_count"] == 1
        assert summary["matched_count"] == 1
        assert summary["issue_count"] == 0
        assert summary["warnings"] == []
        assert summary["rows"][0]["source_sheet"] == "导入数据"
        assert summary["rows"][0]["source_row"] == 4
        assert summary["rows"][0]["match_status"] == "MATCHED"
