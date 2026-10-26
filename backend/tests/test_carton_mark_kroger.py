import copy
import io
import json

import pytest
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image
from PIL import Image as PillowImage
from pypdf import PdfReader

from app.services import carton_mark_kroger as kroger
from app.services import carton_mark_ai_pilot as pilot
from app.services import carton_mark_source_notes as notes


def workbook(extra_note=False):
    book=Workbook();s=book.active;s.title="SM"
    s["A1"]="箱唛四面印刷";s["A13"]="KROGER"
    s["A14"]="PURCHASE ORDER # EXAMPLE"
    s["R3"]='ZURU LOGO居中尽可能放大；G.W >=50LBS时四面顶部需印刷"Dual Lift"'
    s["D24"]="不同客户PO内容替换；尺寸与重量请按实际填写"
    s["A26"]="REMARK"
    s["A27"]="GTIN条码印刷在两正唛右下角；Interleaved 2 of 5（无框）；不小于70mm x 35mm；GW NW尺寸按实际"
    s["A17"]="GROSS WEIGHT: XX LBS";s["A18"]="NET WEIGHT: XX LBS"
    s["A19"]='CARTON MEASUREMENT: H-XX"/W-XX"/D-XX"';s["A22"]="MADE IN CHINA"
    for col,text in enumerate(("Customer PO #","SKU/PLN#","Spec#","Customer Description","Case Pack","CTN","GTIN#"),18): s.cell(25,col,text)
    body="1012345678901"
    gtin=body+str((-sum(int(c)*(3 if i%2==0 else 1) for i,c in enumerate(reversed(body))))%10)
    for row,po,article,count in ((26,"PO-001","00123",9),(27,"PO-002","00456",10)):
        for col,value in enumerate((po,article,"SPEC-1","CURRENT PRODUCT","6",str(count),gtin),18): s.cell(row,col,value)
    if extra_note:s["C35"].comment=Comment("另有特殊加工要求，请确认","source")
    pic=io.BytesIO();PillowImage.new("RGB",(200,100),"black").save(pic,format="PNG")
    for anchor in ("B3","N3"):s.add_image(Image(io.BytesIO(pic.getvalue())),anchor)
    out=io.BytesIO();book.save(out);book.close();return out.getvalue()


def plan(source):
    return dict(template_id="kroger-review-v1",layout_variant="paired",
                rows=[dict(source_id=r["source_id"],fields=r["fields"]) for r in source["rows"]],
                notes=[dict(source_id=n.id,rules=sorted(kroger.rules_for(n)),needs_review=False) for n in source["notes"]])


def test_reads_authoritative_po_rows_preserves_zeros_and_collects_cell_note_sources():
    source=kroger.extract("current.xlsx",workbook())
    assert [r["fields"]["article"] for r in source["rows"]]==["00123","00456"]
    assert [r["fields"]["carton_total"] for r in source["rows"]]==["9","10"]
    assert {n.location for n in source["notes"]}=={"SM!A1","SM!R3","SM!D24","SM!A26","SM!A27"}
    assert source["weights"]["gross"]=="XX LBS" and source["image_count"]==2
    kroger.validate(json.dumps(plan(source)),source)


def test_po_table_layout_is_selected_by_structure_and_keeps_the_original_printed_name():
    from openpyxl import load_workbook
    book=load_workbook(io.BytesIO(workbook()))
    book.active["A13"]="ZUMIEZ INC."
    stream=io.BytesIO();book.save(stream);book.close()
    source=kroger.extract("same-collection.xlsx",stream.getvalue())
    assert source["recipient"]=="ZUMIEZ INC."
    content,_=kroger.render(source,api_verified=False)
    text=PdfReader(io.BytesIO(content)).pages[0].extract_text()
    assert "ZUMIEZ INC." in text and "KROGER" not in text


def test_rejects_changed_fields_mixed_po_rows_lost_leading_zero_and_unread_or_misinterpreted_notes():
    source=kroger.extract("current.xlsx",workbook())
    for change in ("count","zero","missing_row","missing_note","wrong_rule"):
        result=copy.deepcopy(plan(source))
        if change=="count":result["rows"][0]["fields"]["carton_total"]="99"
        if change=="zero":result["rows"][0]["fields"]["article"]="123"
        if change=="missing_row":result["rows"].pop()
        if change=="missing_note":result["notes"].pop()
        if change=="wrong_rule":result["notes"][0]["rules"]=["min_barcode_70x35"]
        with pytest.raises(pilot.PilotError):kroger.validate(json.dumps(result),source)


def test_cell_comments_are_extracted_and_unknown_requirements_cannot_be_silently_accepted():
    source=kroger.extract("current.xlsx",workbook(extra_note=True))
    assert any(n.location=="SM!C35/comment" for n in source["notes"])
    with pytest.raises(pilot.PilotError,match="未支持"):
        kroger.validate(json.dumps(plan(source)),source)


def test_prepare_sends_every_table_row_and_note_without_images_and_renders_once(monkeypatch):
    from app.services import carton_supplier_mark_ai as ai
    source=kroger.extract("current.xlsx",workbook());calls=[]
    monkeypatch.setattr(ai,"provider_configuration",lambda:({},"signature"))
    def request(entries,template,connection,**kwargs):
        calls.append(entries)
        assert {e.id for e in entries}=={r["source_id"] for r in source["rows"]}|{n.id for n in source["notes"]}
        assert all(e.kind in {"current_row","current_note"} for e in entries)
        assert "图片没有提供" in kwargs["prompt"]
        return json.dumps(plan(source)),{}
    monkeypatch.setattr(pilot,"request_plan",request)
    monkeypatch.setattr(kroger,"render",lambda source,**kwargs:(b"pdf",3))
    pdf,pages,warnings=kroger.prepare(dict(provider_signature="signature",kroger_source=source))
    assert pdf==b"pdf" and pages==3 and len(calls)==1 and "嵌入参考图未发送" in warnings[-1]


def test_render_keeps_po_pages_separate_missing_values_explicit_and_notes_off_print_faces():
    source=kroger.extract("current.xlsx",workbook())
    content,pages=kroger.render(source,api_verified=False)
    reader=PdfReader(io.BytesIO(content));assert pages==3
    for index,row in enumerate(source["rows"]):
        text=reader.pages[index].extract_text()
        assert row["fields"]["customer_po"] in text and row["fields"]["article"] in text
        assert source["rows"][1-index]["fields"]["customer_po"] not in text
        assert text.count("GTIN# "+row["fields"]["gtin"])==2
        assert "XX LBS" in text and "Dual Lift" not in text and "SM!A27" not in text
    assert "SM!A27" in reader.pages[-1].extract_text()


def test_current_draft_dispatches_kroger_before_zumiez_and_does_not_call_api_on_bad_overrides(monkeypatch):
    from app.services import carton_supplier_mark_draft as draft
    from app.services import carton_supplier_mark_ai as ai
    monkeypatch.setattr(ai,"provider_configuration",lambda:({},"signature"))
    original=draft.original("current.xlsx",workbook(),"deepseek_zumiez",True,{})
    assert original["ai_profile"]=="kroger" and original["unbound"]
    monkeypatch.setattr(kroger,"prepare",lambda source:(b"kroger",3,[]))
    assert draft.prepare(original,"deepseek_zumiez")== (b"kroger",3,[])
    from fastapi import HTTPException
    with pytest.raises(HTTPException) as exc:
        draft.original("current.xlsx",workbook(),"deepseek_zumiez",True,{"carton_total":"99"})
    assert exc.value.status_code==422
