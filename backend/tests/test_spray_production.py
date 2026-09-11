"""Workshop acceptance uses synthetic quantities only, never production databases."""
import importlib
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal as D
from pathlib import Path
from uuid import uuid4
import sys
import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'spray.db'}")
    monkeypatch.setenv("SEED_ADMIN_PASSWORD", "SprayTestOnly123!")
    monkeypatch.setenv("AUTHZ_MODE", "enforce")
    for name in list(sys.modules):
        if name == "app" or name.startswith("app."):
            del sys.modules[name]
    main = importlib.import_module("app.main")
    with TestClient(main.app) as c:
        r = c.post("/api/auth/login", json={"username": "admin", "password": "SprayTestOnly123!"})
        assert r.status_code == 200, r.text
        yield c


def post(c, action, payload, expected=200, key=None, factory="huaxing"):
    r = c.post(f"/api/spray-production/commands/{action}", json={"factory_id": factory, "operation_id": key or uuid4().hex, **payload})
    assert r.status_code == expected, r.text
    return r.json()


def collection(c, kind, factory="huaxing"):
    r = c.get(f"/api/spray-production/collections/{kind}", params={"factory_id": factory})
    assert r.status_code == 200, r.text
    return r.json()["items"]


def setup_order(c, two_steps=False):
    post(c, "orders", {"document_no": "000935373-1", "customer": "华康B", "customer_factory_id": "huakang-b", "due_date": "2026-09-30",
        "lines": [{"product_no": "000935373-1", "part_name": "头部", "quantity": "50000", "price": ".24", "price_reference": "测试确认报价", "steps": [
            {"name": "白底", "capability": "manual"}] + ([{"name": "眼黑", "capability": "pad"}] if two_steps else [])}]})
    line = collection(c, "lines")[0]
    st = sorted(collection(c, "steps"), key=lambda s: s["sequence"])
    post(c, "prepare", {"line_id": line["id"], "ready_at": "2026-08-01T08:00:00+08:00", "reason": "已领料、调油、调机并完成首件"})
    post(c, "resources", {"name": "手喷工位 01", "capability": "manual", "workers": 2, "machine_rate": "1000", "person_minutes": ".12", "setup_hours": ".5",
                           "calendar": [{"start": "2026-09-11T00:00:00+08:00", "end": "2026-09-30T23:59:59+08:00"}]})
    post(c, "receipts", {"document_no": "IN-001", "business_date": "2026-09-11", "lines": [{"source_line": "1", "line_id": line["id"], "accepted": "4900", "rejected": "100"}]})
    return line, st, collection(c, "resources")[0], collection(c, "batches")[0]


def plan(c, batch, step, res, qty, start="2026-09-11T08:00:00+08:00", rework=False):
    p = post(c, "schedule-preview", {"changes": [{"batch_id": batch["id"], "step_id": step["id"], "resource_id": res["id"], "quantity": str(qty), "workers": 2, "start_at": start, "rework": rework}]})
    p = post(c, "schedule-apply", {"scenario_id": p["id"]})
    task = p["ids"][0]
    post(c, "task-state", {"task_id": task, "action": "start"})
    return task


def test_production_to_returns_settlement_and_idempotency(client):
    c = client
    line, st, res, batch = setup_order(c)
    task = plan(c, batch, st[0], res, 4900)
    payload = {"business_date": "2026-09-11", "shift": "白班", "team": "手喷一组", "lines": [{"task_id": task, "regular_qty": "4700", "overtime_qty": "200", "regular_hours": "4", "overtime_hours": "1",
              "good": "4700", "rework": "150", "scrap": "50", "people": [{"name": "历史组员"}]}]}
    original = post(c, "reports", payload, key="one-report")
    for _ in range(2):
        assert post(c, "reports", payload, key="one-report") == original
    post(c, "reports", {**payload, "team": "changed"}, expected=409, key="one-report")
    assert len(collection(c, "report-lines")) == 1
    assert collection(c, "report-lines")[0]["wage_status"] == "unpriced"
    rework = plan(c, batch, st[0], res, 150, "2026-09-11T15:00:00+08:00", True)
    post(c, "reports", {**payload, "lines": [{"task_id": rework, "regular_qty": "150", "good": "120", "scrap": "30", "people": [{"name": "历史组员"}]}]})
    trace = c.get(f"/api/spray-production/batches/{batch['id']}/trace", params={"factory_id": "huaxing"}).json()
    assert D(trace["balances"]["finished"]) == 4820
    assert D(trace["balances"]["scrap"]) == 80
    assert D(trace["balances"]["rejected"]) == 100
    assert sum(D(v) for v in trace["balances"].values()) == 5000
    assert sum(D(r["regular_qty"]) + D(r["overtime_qty"]) for r in collection(c, "report-lines")) == 5050
    post(c, "shipments", {"document_no": "SHIP-01", "customer": "华康B", "business_date": "2026-09-11", "lines": [{"batch_id": batch["id"], "quantity": "1000"}],
                         "containers": [{"name": "胶箱", "quantity": "10", "unit": "个"}]})
    shipment = collection(c, "shipment-lines")[0]
    post(c, "returns", {"shipment_line_id": shipment["id"], "quantity": "100", "document_no": "RETURN-01", "business_date": "2026-09-12", "reason": "客户退回待判"})
    preview = c.post("/api/spray-production/settlements/preview", json={"factory_id": "huaxing", "customer": "华康B", "period": "2026-09", "currency": "HKD"}).json()
    assert D(preview["amount"]) == 216
    post(c, "settlements", {**preview, "line_ids": sorted(v["id"] for v in preview["lines"])})
    post(c, "returns", {"shipment_line_id": shipment["id"], "quantity": "1", "document_no": "RETURN-02", "business_date": "2026-09-13", "reason": "已冻结"}, expected=409)
    post(c, "returns", {"shipment_line_id": shipment["id"], "quantity": "20", "document_no": "RETURN-02", "business_date": "2026-10-01", "reason": "下期贷项"})
    assert D(collection(c, "settlements")[0]["amount"]) == 216
    assert not collection(c, "orders", "huakang-b")
    assert c.get(f"/api/spray-production/batches/{batch['id']}/trace", params={"factory_id": "huakang-b"}).status_code == 404
    assert c.get("/api/spray-production/summary", params={"factory_id": "group"}).status_code == 422


def test_no_forecast_stock_no_downstream_overconsumption_atomic_report(client):
    c = client
    line, st, res, batch = setup_order(c, True)
    change = {"batch_id": batch["id"], "step_id": st[0]["id"], "resource_id": res["id"], "quantity": "50000", "workers": 2, "start_at": "2026-09-11T08:00:00+08:00"}
    post(c, "schedule-preview", {"changes": [change]}, expected=409)
    task = plan(c, batch, st[0], res, 2000)
    good = {"task_id": task, "regular_qty": "2000", "good": "2000", "people": [{"name": "组员"}]}
    post(c, "reports", {"business_date": "2026-09-11", "shift": "白班", "team": "组", "lines": [good, good]}, expected=422)
    assert not collection(c, "reports")
    post(c, "reports", {"business_date": "2026-09-11", "shift": "白班", "team": "组", "lines": [good]})
    trace = c.get(f"/api/spray-production/batches/{batch['id']}/trace", params={"factory_id": "huaxing"}).json()
    assert D(trace["balances"]["ready:" + st[1]["id"]]) == 2000
    assert "finished" not in trace["balances"]


def test_material_actual_cost_and_confirmed_rate_versions(client):
    c = client
    p = post(c, "purchases", {"document_no": "FBD 0112802", "supplier": "倍福", "material": "628A软胶开油水", "quantity": "2", "unit": "桶", "base_unit": "KG", "conversion": "160", "conversion_evidence": "原采购单 160KG/桶", "price": "1472", "currency": "CNY"})
    for kind, qty in [("receipt", "320"), ("issue", "30"), ("consume", "20"), ("return", "10")]:
        post(c, "material-events", {"purchase_id": p["id"], "kind": kind, "quantity": qty, "business_date": "2026-09-11", "document_no": kind})
    events = collection(c, "material-events")
    assert sum(D(e["cost"]) for e in events) == D(184)
    post(c, "material-events", {"purchase_id": p["id"], "kind": "consume", "quantity": "1", "business_date": "2026-09-11", "document_no": "over"}, expected=409)
    rule = post(c, "rates", {"name": "旧表复算", "rule_code": "legacy_piece_normalized", "currency": "HKD", "quantity_basis": "attempt", "effective_date": "2026-01-01", "parameters": {"rate": ".82", "base_hours": "11", "paid_hours": "12.5", "rmb_per_hkd": ".88"}, "evidence": "S09"})
    post(c, "rate-activate", {"rate_id": rule["id"], "reason": "不允许默认套用"}, expected=422)
    service = importlib.import_module("app.services.spray_production")
    assert sum(D(v) for v in service.allocation(D("100"), [D(1)] * 3)) == 100
    assert D(service.throughput(D(5000), D(1000), D(".12"), 2, D(".5"))["hours"]) == D("5.5")


def test_painting_clerk_is_scoped_and_costs_never_downloaded(client):
    c = client
    setup_order(c)
    auth = importlib.import_module('app.services.auth')
    model = importlib.import_module('app.models.auth')
    dbmod = importlib.import_module('app.db')
    with dbmod.SessionLocal() as db:
        salt, hashed = auth.make_password_hash('ClerkTestOnly123!')
        db.add(model.AuthUser(id='spray-clerk', username='spray-clerk', display_name='喷油测试文员', password_salt=salt, password_hash=hashed, status='active', force_password_change=0, created_at=auth.now_text(), updated_at=auth.now_text()))
        db.add(model.AuthUserRole(id='spray-clerk-role', user_id='spray-clerk', role_id='position_painting_clerk', factory_id='huaxing', department='production'))
        db.commit()
    assert c.post('/api/auth/login', json={'username':'spray-clerk','password':'ClerkTestOnly123!'}).status_code == 200
    assert c.get('/api/spray-production/summary', params={'factory_id':'huaxing'}).status_code == 200
    assert c.get('/api/spray-production/summary', params={'factory_id':'huakang-b'}).status_code == 403
    lines = collection(c,'lines')
    assert 'price' not in lines[0] and 'currency' not in lines[0]
    assert c.get('/api/spray-production/collections/rates', params={'factory_id':'huaxing'}).status_code == 403
    assert c.get('/api/spray-production/imports/nonexistent', params={'factory_id':'huaxing'}).status_code == 403
    assert c.get('/api/spray-production/exports/wages', params={'factory_id':'huaxing'}).status_code == 403
    post(c, 'reports', {'lines':['invalid']}, expected=422)
    post(c, 'orders', {'lines':[{'price':'.1'}]}, expected=403)


def test_suggestions_quality_corrections_and_exports(client):
    c=client
    line,st,res,batch=setup_order(c)
    proposal=post(c,'schedule-suggest',{'start_at':'2026-09-11T08:00:00+08:00'})
    assert len(proposal['changes'])==1 and D(proposal['changes'][0]['quantity'])==4900
    task=post(c,'schedule-apply',{'scenario_id':proposal['id']})['ids'][0]
    post(c,'task-state',{'task_id':task,'action':'start'})
    data={'business_date':'2026-09-11','shift':'白班','team':'测试','lines':[{'task_id':task,'regular_qty':'1000','good':'1000','people':[{'name':'测试'}]}]}
    original=post(c,'reports',data)
    post(c,'shipments',{'document_no':'SH-01','customer':'华康B','business_date':'2026-09-11','lines':[{'batch_id':batch['id'],'quantity':'100'}]})
    shipped=collection(c,'shipment-lines')[0]
    post(c,'returns',{'document_no':'RT-01','shipment_line_id':shipped['id'],'business_date':'2026-09-11','quantity':'100','reason':'测试'})
    post(c,'quality',{'batch_id':batch['id'],'source_state':'return-held','disposition':'release','quantity':'100','reason':'实检放行'})
    # Restored total balance is not permission to rewrite already-consumed lineage.
    post(c,'report-correct',{**data,'report_id':original['id'],'reason':'已有下游流转'},expected=409)
    post(c,'quality',{'batch_id':batch['id'],'source_state':'receipt-held','disposition':'rework','quantity':'1','reason':'无效组合'},expected=422)
    from io import BytesIO
    from openpyxl import load_workbook
    response=c.get('/api/spray-production/exports/orders',params={'factory_id':'huaxing'})
    assert response.status_code==200
    book=load_workbook(BytesIO(response.content))
    assert book.active['A3'].value=='000935373-1'
    other=c.get('/api/spray-production/exports/orders',params={'factory_id':'huakang-b'})
    assert load_workbook(BytesIO(other.content)).active.max_row==2


def test_price_versions_partial_credit_rounding_and_input_validation(client):
    c=client
    line,st,res,batch=setup_order(c)
    post(c,'price-set',{'line_id':line['id'],'expected_revision':0,'price':'.335','currency':'HKD','reason':'stale'},expected=409)
    current=collection(c,'lines')[0]
    post(c,'price-set',{'line_id':line['id'],'expected_revision':current['revision'],'price':'.335','currency':'HKD','reason':'确认测试单价'})
    post(c,'resources',{'name':'invalid','capability':'manual','workers':'.5'},expected=422)
    task=plan(c,batch,st[0],res,3)
    data={'business_date':'2026-09-11','shift':'白班','team':'测试','lines':[{'task_id':task,'regular_qty':'3','good':'3','people':[{'name':'测试'}]}]}
    post(c,'reports',{**data,'lines':[{**data['lines'][0],'people':[123]}]},expected=422)
    post(c,'reports',data)
    post(c,'shipments',{'document_no':'CENT-01','customer':'华康B','business_date':'2026-09-11','lines':[{'batch_id':batch['id'],'quantity':'3'}]})
    shipped=collection(c,'shipment-lines')[0]
    assert D(shipped['amount'])==D('1.01')
    current=collection(c,'lines')[0]
    post(c,'price-set',{'line_id':line['id'],'expected_revision':current['revision'],'price':'.9','currency':'HKD','reason':'后续新价'})
    for index in range(3):
        post(c,'returns',{'document_no':f'CENT-RT-{index}','shipment_line_id':shipped['id'],'business_date':'2026-09-11','quantity':'1','reason':'分次退回'})
    assert sum(D(row['amount']) for row in collection(c,'shipment-lines'))==0
    assert all(D(row['price'])==D('.335') for row in collection(c,'shipment-lines'))


def test_golden_reservations_wage_idempotency_and_exact_price_snapshots(client):
    c=client
    line,st,res,_=setup_order(c)
    batch_id=post(c,'receipts',{'document_no':'GOLDEN-5000','business_date':'2026-09-11','lines':[{'line_id':line['id'],'source_line':'1','accepted':'5000'}]})['id']
    change={'batch_id':batch_id,'step_id':st[0]['id'],'resource_id':res['id'],'quantity':'2000','workers':2,'start_at':'2026-09-11T08:00:00+08:00'}
    scenario=post(c,'schedule-preview',{'changes':[change]})
    task=post(c,'schedule-apply',{'scenario_id':scenario['id']})['ids'][0]
    balances=c.get('/api/spray-production/balances',params={'factory_id':'huaxing'}).json()[batch_id]
    assert D(balances['ready:'+st[0]['id']])==3000
    post(c,'schedule-preview',{'changes':[{**change,'quantity':'3001','start_at':'2026-09-11T15:00:00+08:00'}]},expected=409)
    rate=post(c,'rates',{'name':'待核计件测试','rule_code':'piece_direct','quantity_basis':'good','currency':'CNY','effective_date':'2026-01-01','parameters':{'rate':'.02'},'evidence':'仅隔离测试，不是正式规则'})
    post(c,'task-state',{'task_id':task,'action':'start'})
    report={'business_date':'2026-09-11','shift':'白班','team':'测试','lines':[{'task_id':task,'regular_qty':'2000','good':'2000','rate_id':rate['id'],'people':[{'name':'测试'}]}]}
    for _ in range(3): post(c,'reports',report,key='demo-report-001')
    saved=collection(c,'report-lines');assert len(saved)==1 and saved[0]['wage_status']=='provisional'
    assert D(saved[0]['calculation']['amount'])==40
    for index,price in enumerate(('.24','.26')):
        if index:
            current=collection(c,'lines')[0]
            post(c,'price-set',{'line_id':line['id'],'expected_revision':current['revision'],'price':price,'currency':'HKD','reason':'黄金新价版本'})
        post(c,'shipments',{'document_no':f'GOLDEN-PRICE-{index}','customer':'华康B','business_date':'2026-09-11','lines':[{'batch_id':batch_id,'quantity':'1000'}]})
    assert sorted(D(row['amount']) for row in collection(c,'shipment-lines'))==[D(240),D(260)]
