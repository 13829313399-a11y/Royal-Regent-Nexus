"""Manual carton headers must not shift or populate any business data."""
from copy import copy
from io import BytesIO
from hashlib import sha256
import os
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from app.services import customer_order_unified as service
from app.services.customer_order_360_carton import SHEETS, ensure_carton_columns
from app.services.customer_order_360_workbook import organize_workbook
from test_customer_order_360_version import kwargs, parse_stub, workbook_bytes
from test_customer_order_huakang_a_unified import template


def apply(workbook, factory="huakang-a", customer="360"):
    return ensure_carton_columns(workbook, factory_id=factory, customer_code=customer)


def snapshot(workbook):
    return {s.title: {(r, c): ((cell.value.text, cell.value.ref)
                             if cell.data_type == "f" and not isinstance(cell.value, str)
                             else cell.value, copy(cell._style))
                     for (r, c), cell in s._cells.items()} for s in workbook}


def test_append_preserves_every_existing_cell_formula_and_header_style():
    workbook = openpyxl.load_workbook(BytesIO(template("360")))
    for sheet in workbook:
        sheet["AS3"] = "人工备注"
        sheet["AS3"].font = Font(name="Arial", bold=True, color="FFFFFF")
        sheet["AS3"].fill = PatternFill("solid", fgColor="123456")
        sheet["AT4"] = "不能覆盖"
        sheet["AU4"] = "=1+2"
    before = snapshot(workbook)
    result = apply(workbook)
    after = snapshot(workbook)
    assert set(result["added"]) == set(SHEETS)
    for sheet in workbook:
        assert sheet["AV3"].value == "纸箱"
        assert copy(sheet["AV3"].font) == copy(sheet["AS3"].font)
        assert copy(sheet["AV3"].fill) == copy(sheet["AS3"].fill)
        assert sheet["AV3"].alignment.wrap_text
        assert sheet.column_dimensions["AV"].width == 18
        for position, expected in before[sheet.title].items():
            cell = sheet._cells[position]
            assert cell._style == expected[1]
            assert after[sheet.title][position] == expected
        assert not any(c.value is not None for c in sheet._cells.values()
                       if c.column == 48 and c.row != 3)
    second = apply(workbook)
    assert second["added"] == [] and second["columns"] == result["columns"]


def test_existing_header_aliases_duplicates_values_and_formulas_are_not_rewritten():
    workbook = openpyxl.load_workbook(BytesIO(template("360")))
    for sheet in workbook:
        sheet["AS3"], sheet["AT3"] = " 紙 箱\n", "纸箱"
        sheet["AS4"], sheet["AT4"] = "已下单", "=1+2"
    before = workbook_bytes(workbook)
    restored = openpyxl.load_workbook(BytesIO(before))
    assert apply(restored)["added"] == []
    after = openpyxl.load_workbook(BytesIO(workbook_bytes(restored)))
    for sheet in after:
        assert sheet["AS3"].value == " 紙 箱\n"
        assert sheet["AS4"].value == "已下单" and sheet["AT4"].value == "=1+2"
        assert sheet.max_column == 46


def test_pricing_or_packing_headers_are_not_mistaken_for_carton_header():
    workbook = openpyxl.load_workbook(BytesIO(template("360")))
    for sheet in workbook:
        sheet["AS3"] = "装好纸箱报库存价HK$"
    result = apply(workbook)
    assert set(result["added"]) == set(SHEETS)
    assert all(workbook[n]["AT3"].value == "纸箱" for n in SHEETS)


def test_distant_formatting_does_not_create_a_gap_before_carton_header():
    workbook = openpyxl.load_workbook(BytesIO(template("360")))
    for sheet in workbook:
        sheet["XFD4"].font = Font(italic=True)
    result = apply(workbook)
    assert result["columns"]["接单表"] == ["U"]
    assert result["columns"]["正单评审表"] == ["U"]
    assert all(sheet["XFD4"].font.italic for sheet in workbook)


def test_append_skips_uninstantiated_validation_and_table_ranges():
    workbook = openpyxl.load_workbook(BytesIO(template("360")))
    item = workbook["ITEM表"]
    rule = DataValidation(type="list", formula1='"是,否"')
    rule.add("AZ4:AZ12")
    item.add_data_validation(rule)
    summary = workbook["接单表"]
    summary.add_table(Table(displayName="Existing", ref="BA3:BB4"))
    assert apply(workbook)["columns"] == {"ITEM表": ["BA"], "接单表": ["BC"], "正单评审表": ["U"]}
    assert rule.formula1 == '"是,否"' and str(rule.sqref) == "AZ4:AZ12"


@pytest.mark.parametrize("factory,customer", [("huaxing", "360"), ("huakang-a", "green-toys"), ("huakang-d", "360")])
def test_other_factory_or_customer_is_noop(factory, customer):
    workbook = openpyxl.load_workbook(BytesIO(template("360")))
    before = snapshot(workbook)
    assert not apply(workbook, factory, customer)["applied"]
    assert snapshot(workbook) == before
    assert not any(c.value == "纸箱" for s in workbook for c in s._cells.values())


def test_full_width_fails_before_changing_any_sheet():
    workbook = openpyxl.load_workbook(BytesIO(template("360")))
    workbook["正单评审表"]["XFD4"] = "保留"
    with pytest.raises(service.CustomerOrderUnifiedError, match="无法安全追加"):
        apply(workbook)
    assert not any(c.value == "纸箱" for s in workbook for c in s._cells.values())


def test_real_export_and_standalone_organizer_roundtrip_preserve_manual_carton(monkeypatch):
    parse_stub(monkeypatch)
    output, _, _ = service.export_unified_customer_schedule(**kwargs(template("360")))
    workbook = openpyxl.load_workbook(BytesIO(output))
    for sheet in workbook:
        column, = [c.column for c in sheet[3] if c.value == "纸箱"]
        reference_column = 5 if sheet.title == "ITEM表" else 4
        row = next(c.row for c in sheet._cells.values()
                   if c.column == reference_column and c.value == "RL-765431-23")
        assert sheet.cell(row, column).value is None
        sheet.cell(row, column, "纸箱已下单，待到货")
    first, _, report, _ = organize_workbook("schedule.xlsx", workbook_bytes(workbook))
    assert report["carton_columns"]["added"] == []
    second, _, _, _ = organize_workbook("schedule.xlsx", first)
    result = openpyxl.load_workbook(BytesIO(second))
    for sheet in result:
        column, = [c.column for c in sheet[3] if c.value == "纸箱"]
        assert sum(c.value == "纸箱已下单，待到货" for c in sheet._cells.values() if c.column == column) == 1


def test_private_schedule_changes_only_three_headers_vs_existing_organizer(monkeypatch, tmp_path):
    """Opt-in acceptance against a private workbook, never stored in Git."""
    configured = os.environ.get("CUSTOMER_ORDER_360_CARTON_SAMPLE")
    if not configured:
        pytest.skip("Set CUSTOMER_ORDER_360_CARTON_SAMPLE to the private unified schedule")
    path = Path(configured)
    source = path.read_bytes()
    from app.services import customer_order_360_carton as carton
    with monkeypatch.context() as context:
        context.setattr(carton, "ensure_carton_columns", lambda *args, **kwargs: {})
        baseline, _, _, _ = organize_workbook(path.name, source)
    output, _, report, _ = organize_workbook(path.name, source)
    before = openpyxl.load_workbook(BytesIO(baseline))
    after = openpyxl.load_workbook(BytesIO(output))
    assert before.sheetnames == after.sheetnames
    before_cells = snapshot(before)
    after_cells = snapshot(after)
    checked = 0
    for name in before.sheetnames:
        carton_cells = {(3, openpyxl.utils.column_index_from_string(c))
                        for c in report["carton_columns"]["columns"].get(name, [])}
        for position, expected in before_cells[name].items():
            if position in carton_cells:
                assert expected[0] is None  # Only a previously blank header may acquire a label/style.
                continue
            assert after_cells[name][position] == expected, (name, position)
            checked += 1
        changed = {p for p, value in after_cells[name].items() if before_cells[name].get(p) != value}
        if name in SHEETS:
            assert changed == carton_cells
            row, column = next(iter(changed))
            assert row == 3 and after[name].cell(row, column).value == "纸箱"
        else:
            assert not changed
        assert str(before[name].merged_cells) == str(after[name].merged_cells)
        assert before[name].data_validations == after[name].data_validations
        assert before[name].print_area == after[name].print_area
        assert before[name].freeze_panes == after[name].freeze_panes
    assert sha256(path.read_bytes()).digest() == sha256(source).digest()
    output_path = tmp_path / "360-carton-acceptance.xlsx"
    output_path.write_bytes(output)
    print({"checked_existing_cells": checked, "carton_columns": report["carton_columns"],
           "output_for_visual_check": str(output_path)})
