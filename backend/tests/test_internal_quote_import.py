from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as WorksheetImage
from PIL import Image as PillowImage

from app.services.internal_quote_artifacts import _merge_import_payload
from app.services.internal_quote_import import find_header, parse_internal_quote_workbook, workbook_rows
from app.services.internal_quote_templates import (
    FIXED_TEMPLATE_FILE_NAMES,
    TEMPLATE_LABELS,
    build_internal_quote_import_template,
)


def workbook_bytes(rows: list[list[object]], title: str = "报价明细") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.mark.parametrize("import_type", sorted(TEMPLATE_LABELS))
def test_downloadable_import_template_uses_a_header_recognized_by_its_parser(import_type: str):
    content, file_name = build_internal_quote_import_template(import_type)

    if import_type == "hair":
        assert content.startswith(b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1")
        assert file_name.endswith(".xls")
    else:
        assert content.startswith(b"PK")
        assert file_name.endswith(".xlsx")
    if import_type in FIXED_TEMPLATE_FILE_NAMES:
        assert file_name == FIXED_TEMPLATE_FILE_NAMES[import_type]
        expected_path = (
            Path(__file__).resolve().parents[1]
            / "app"
            / "data"
            / "internal_quote_import_templates"
            / file_name
        )
        assert content == expected_path.read_bytes()
        parsed = parse_internal_quote_workbook(content, import_type)
        if import_type == "mold":
            assert parsed.sheet_name == "01"
            assert parsed.header_row == 11
            assert parsed.row_count == 0
            workbook = load_workbook(BytesIO(content), data_only=False, read_only=True)
            try:
                assert [
                    workbook["01"].cell(11, column).value
                    for column in range(10, 14)
                ] == ["产能", "机型", "胶件重量（g)", "Remark备注"]
            finally:
                workbook.close()
        else:
            assert parsed.row_count > 0
        return

    sheet_name, _rows, header_index = find_header(workbook_rows(content), import_type)
    assert sheet_name == TEMPLATE_LABELS[import_type]
    assert header_index == 0


def test_downloadable_sewing_template_is_the_exchange_rate_reference_workbook():
    content, file_name = build_internal_quote_import_template("sewing")
    workbook = load_workbook(BytesIO(content), data_only=False, read_only=True)
    try:
        assert file_name == "车缝报价单.xlsx"
        assert workbook.sheetnames == ["总表260707", "明细260707"]
        detail = workbook["明细260707"]
        assert [detail.cell(1, column).value for column in range(1, 13)] == [
            "物料名称", "裁片部位", "供应商", "布料MOQ/Y", "低于MOQ/每色产生费用",
            "用量/码", "单价", "汇率", "成本", "码点", "价钱", "备注",
        ]
        assert detail["I3"].value == "=F3*G3/H3"
        assert detail["K3"].value == "=I3*J3"
        assert detail["K28"].value == "=SUM(K3:K27)"
    finally:
        workbook.close()


def water_table_workbook_with_image() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "模价表1"
    sheet.append(["NO.", "MOLD NO.", "Tool", "PART NAME", None, "CAV", "UP", "MAT'L", "COLOR", "CLAMPING FORCE (tonne)", "Part Weight (g)", "TOOL INFORMATION", "Mold Weight (kg)", "Tool Insert Mat'l", "No. of Slide", "CYCLE", "DAILY", "WEEKLY", None, "TOTAL AMOUNT", "PICTURES", "REMARKS"])
    sheet.append([None, None, "Type", "DESCRIPTION", "CHINESE NAME", None, None, None, None, None, None, "Dim (HxWxD cm)", None, None, None, "TIME (sec)", "RATE", "RATE(K)", "GATE", "(RMB)", None, None])
    sheet.append([1, "M01", "INJ", None, "水桌主体", 1, 1, "PP", "Blue", "900T", None, "90*90*80", None, "718H", None, 95, None, 4752, "热流道", 237000, None, None])
    image_bytes = BytesIO()
    PillowImage.new("RGB", (4, 4), color=(16, 118, 110)).save(image_bytes, format="PNG")
    image_bytes.seek(0)
    image = WorksheetImage(image_bytes)
    image.anchor = "U3"
    sheet.add_image(image)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_mold_import_maps_rr2_fields_to_p2_engineering_contract():
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                ["模号", "产品名称", "材质", "克重", "套数", "模价", "机型", "目标数", "备注"],
                ["M-100", "公仔头模", "ABS", 85, 1, 28000, "4A", 8000, "客户模"],
            ]
        ),
        "mold",
        fallback_qty=Decimal("5000"),
    )

    assert parsed.target_department == "engineering"
    assert parsed.header_row == 1
    assert parsed.row_count == 1
    assert parsed.payload_fragment["amortization_qty"] == "5000.0000"
    assert parsed.payload_fragment["molds"][0] == {
        "item": "公仔头模",
        "mold_no": "M-100",
        "chinese_name": "",
        "quantity": "1.0000",
        "cost_rmb": "28000.0000",
        "material": "ABS",
        "material_type": "",
        "net_weight_g": "85.0000",
        "cavity": "",
        "machine_code": "4A",
        "target_output": "8000.0000",
        "mold_base_type": "",
        "mold_base_material": "",
        "structure": "",
        "process": "",
        "cycle_time_seconds": "0.0000",
        "mold_size": "",
        "mold_specification": "",
        "color": "",
        "image_reference": "",
        "image_attachment_ids": [],
        "remark": "客户模",
        "source_row": 2,
    }
    assert any("图片附件" in warning for warning in parsed.warnings)


def test_mold_import_maps_water_table_two_row_headers_and_skips_blank_rows_after_m12():
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                ["NO.", "MOLD NO.", "Tool", "PART NAME", None, "CAV", "UP", "MAT'L", "COLOR", "CLAMPING FORCE (tonne)", "Part Weight (g)", "TOOL INFORMATION", "Mold Weight (kg)", "Tool Insert Mat'l", "No. of Slide", "CYCLE", "DAILY", "WEEKLY", None, "TOTAL AMOUNT", "PICTURES", "REMARKS"],
                [None, None, "Type", "DESCRIPTION", "CHINESE NAME", None, None, None, None, None, None, "Dim (HxWxD cm)", None, None, None, "TIME (sec)", "RATE", "RATE(K)", "GATE", "(RMB)", None, None],
                [11, "M11", "INJ", None, "提手", 8, 2.67, "PP", "Blue", "160T", None, "35*55*45", None, "718H", "4个行位", 35, None, 34438, "细水口", 44000, None, None],
                [12, "M12", "INJ", None, "水桶1/水桶2/滚筒/挂钩/起动臂左/起动臂右", 6, 1, "PP", "021C Orange", "160T", None, "40*55*44", None, "718H", None, 35, None, 12898, "潜水口", 49000, None, None],
                [],
                [None] * 22,
                ["备注：付款方式"],
                [None] * 18 + ["合计(RMB)", 93000],
            ]
        ),
        "mold",
    )

    assert parsed.row_count == 2
    assert [row["mold_no"] for row in parsed.payload_fragment["molds"]] == ["M11", "M12"]
    m12 = parsed.payload_fragment["molds"][1]
    assert m12["item"] == "水桶1/水桶2/滚筒/挂钩/起动臂左/起动臂右"
    assert m12["chinese_name"] == ""
    assert m12["material_type"] == "PP"
    assert m12["mold_base_material"] == "718H"
    assert m12["process"] == "潜水口"
    assert m12["mold_size"] == "40*55*44"
    assert m12["mold_specification"] == ""
    assert m12["image_reference"] == ""
    assert m12["image_attachment_ids"] == []
    assert m12["cost_rmb"] == "49000.0000"
    assert any("U 列未识别" in warning for warning in parsed.warnings)


def test_mold_import_extracts_embedded_picture_bytes_and_maps_them_to_source_row():
    parsed = parse_internal_quote_workbook(water_table_workbook_with_image(), "mold")

    assert len(parsed.embedded_images) == 1
    assert parsed.embedded_images[0].source_row == 3
    assert parsed.embedded_images[0].content.startswith(b"\x89PNG\r\n\x1a\n")
    assert parsed.payload_fragment["molds"][0]["image_reference"] == "模具图片-U3-1.png"
    assert parsed.payload_fragment["molds"][0]["image_attachment_ids"] == []
    assert any("已识别并提取 U 列 1 张嵌入图片" in warning for warning in parsed.warnings)


def zhanxing_workbook_with_image() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "01"
    for _ in range(10):
        sheet.append([])
    sheet.append([
        "Item No. 模号",
        "Item Description 项目内容",
        "Material 原料",
        "Cavities 件",
        "Up套",
        "图片",
        "Mould Size(L*W*H) 工模尺寸",
        "Mould Prices 模价(RMB)",
        "钢材/钢料硬度",
        "产能",
        "机型",
        "胶件重量（g)",
        "Remark备注",
    ])
    sheet.append([
        "M-01", "水桌主体", "PP", 1, 1, None,
        "90*90*80", 237000, "718H", 3000, "150T", 32, "热流道",
    ])
    image_bytes = BytesIO()
    PillowImage.new("RGB", (4, 4), color=(16, 118, 110)).save(image_bytes, format="PNG")
    image_bytes.seek(0)
    image = WorksheetImage(image_bytes)
    image.anchor = "F12"
    sheet.add_image(image)
    sheet.append(["M-02", "顶部桌面", "PP", 1, 1, None, "60*60*56", 100000, "718H", 3000, "150T", 25, "细水口"])
    sheet.append(["M-03", "腿", "PP", 2, .67, None, "65*65*56", 103000, "718H", 3000, "60-80T", 17, ""])
    sheet.append([])
    sheet.append(["工模 ( 注塑模) - 套"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_mold_import_maps_zhanxing_fixed_template_columns_and_embedded_images():
    parsed = parse_internal_quote_workbook(
        zhanxing_workbook_with_image(),
        "mold",
        fallback_qty=Decimal("3000"),
    )

    assert parsed.target_department == "engineering"
    assert parsed.sheet_name == "01"
    assert parsed.header_row == 11
    assert parsed.row_count == 3
    assert parsed.payload_fragment["amortization_qty"] == "3000.0000"
    assert [
        (
            row["item"],
            row["mold_no"],
            row["chinese_name"],
            row["material_type"],
            row["cavity"],
            row["quantity"],
            row["mold_size"],
            row["cost_rmb"],
            row["mold_base_material"],
            row["target_output"],
            row["machine_code"],
            row["net_weight_g"],
            row["image_reference"],
        )
        for row in parsed.payload_fragment["molds"]
    ] == [
        ("水桌主体", "M-01", "水桌主体", "PP", "1", "1.0000", "90*90*80", "237000.0000", "718H", "3000.0000", "150T", "32.0000", "模具图片-F12-1.png"),
        ("顶部桌面", "M-02", "顶部桌面", "PP", "1", "1.0000", "60*60*56", "100000.0000", "718H", "3000.0000", "150T", "25.0000", ""),
        ("腿", "M-03", "腿", "PP", "2", "0.6700", "65*65*56", "103000.0000", "718H", "3000.0000", "60-80T", "17.0000", ""),
    ]
    assert [
        (image.source_row, image.file_name)
        for image in parsed.embedded_images
    ] == [
        (12, "模具图片-F12-1.png"),
    ]
    assert all(image.content.startswith(b"\x89PNG\r\n\x1a\n") for image in parsed.embedded_images)
    assert any("已识别并提取 F 列 1 张嵌入图片" in warning for warning in parsed.warnings)


def test_mold_import_keeps_legacy_zhanxing_combined_capacity_machine_column():
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                [
                    "Item No. 模号",
                    "Item Description 项目内容",
                    "Material 原料",
                    "Cavities 件",
                    "Up套",
                    "图片",
                    "Mould Size(L*W*H) 工模尺寸",
                    "Mould Prices 模价(RMB)",
                    "钢材/钢料硬度",
                    "产能/机型",
                    "胶件重量（g)",
                    "Remark备注",
                ],
                [
                    "M-OLD",
                    "旧版模具",
                    "ABS",
                    2,
                    1,
                    None,
                    "40*50*60",
                    28000,
                    "718H",
                    "3600/120T",
                    18,
                    "旧版兼容",
                ],
            ],
            title="01",
        ),
        "mold",
        fallback_qty=Decimal("3000"),
    )

    row = parsed.payload_fragment["molds"][0]
    assert row["target_output"] == "3600.0000"
    assert row["machine_code"] == "120T"
    assert row["net_weight_g"] == "18.0000"
    assert row["remark"] == "旧版兼容"


def test_electronic_import_converts_rmb_and_maps_extra_parameters():
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                ["零件名称", "规格", "用量", "单价RMB", "备注"],
                ["IC", "A1", 2, 10, "主控"],
                ["邦定成本", 1.5],
                ["10%利润"],
            ]
        ),
        "electronic",
        rmb_hkd=Decimal("0.85"),
    )

    component = parsed.payload_fragment["components"][0]
    assert parsed.target_department == "electronic"
    assert component["quantity"] == "2.0000"
    assert component["unit_price_rmb"] == "10.0000"
    assert component["unit_price_hkd"] == "11.7647"
    assert component["tax_rate_percent"] == "13.0000"
    assert component["source_currency"] == "RMB"
    assert parsed.payload_fragment["bonding_rmb"] == "1.5000"
    assert parsed.payload_fragment["bonding_hkd"] == "1.7647"
    assert parsed.payload_fragment["profit_rate_percent"] == "10.0000"


def test_electronic_import_keeps_smt_specs_as_components_and_reads_exact_summary_labels():
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                ["零件名称", "规格", "用量", "单价RMB", "合计RMB", "备注"],
                ["IC", "A1", 1, 0.1, 0.1, "主控"],
                [None, "A2", 2, 0.2, 0.4, ""],
                ["电阻", "SMT0603", 18, 0.002, 0.036, ""],
                [None, "SMT0402", 1, 0.003, 0.003, ""],
                ["报价按客户确认图纸为准"],
                [None, None, None, "零件成本:", 0.539],
                [None, None, None, "贴片成本:", 0.496],
                [None, None, None, "人工成本:", 0.6925],
                ["含*10%利润价(月结以此价为准)"],
            ]
        ),
        "electronic",
        rmb_hkd=Decimal("0.85"),
    )

    assert parsed.row_count == 4
    assert len(parsed.payload_fragment["components"]) == 2
    assert parsed.payload_fragment["components"][1]["specification"] == "SMT0603"
    assert parsed.payload_fragment["components"][1]["children"][0]["specification"] == "SMT0402"
    assert parsed.payload_fragment["smt_rmb"] == "0.4960"
    assert parsed.payload_fragment["smt_hkd"] == "0.5835"
    assert parsed.payload_fragment["labor_rmb"] == "0.6925"
    assert parsed.payload_fragment["profit_rate_percent"] == "10.0000"


def test_electronic_import_replaces_complete_quote_even_when_append_is_requested():
    merged = _merge_import_payload(
        "electronic",
        {
            "components": [{"item": "旧件", "quantity": "1", "unit_price_hkd": "2"}],
            "bonding_hkd": "1",
            "smt_hkd": "2",
            "profit_rate_percent": "10",
            "tax_credit_difference_hkd": "9",
        },
        {
            "pricing_currency": "RMB",
            "components": [{"item": "新件", "quantity": "1", "unit_price_rmb": ".3"}],
            "bonding_rmb": ".5",
            "bonding_hkd": ".5882",
            "smt_rmb": ".2",
            "smt_hkd": ".2353",
            "profit_rate_percent": "10",
        },
        "append",
        ".85",
    )

    assert len(merged["components"]) == 1
    assert merged["components"][0]["item"] == "新件"
    assert merged["pricing_currency"] == "RMB"
    assert merged["bonding_rmb"] == ".5"
    assert merged["smt_rmb"] == ".2"
    assert "bonding_hkd" not in merged
    assert "smt_hkd" not in merged
    assert "tax_credit_difference_hkd" not in merged


def test_molding_import_maps_rr2_injection_and_blow_sections_and_ignores_derived_prices():
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                ["二、注塑部分"],
                ["模具名称", "模号", "材质", "料型", "颜色", "啤净重(g)", "料损耗 3%", "料价 HK$/g", "原料单价 HK$", "机台", "啤价(HK$/啤)", "出模数", "套数", "机型", "目标数", "周期(秒)", "成品金额 HK$", "成品用量", "备注"],
                ["主体模", "M-01", "ABS", "750SW", "黑色", 100, 103, 0.0187, 1.926, "80T", 0.188, "2", 1, "5A", 5000, 24, 2.114, 2, "客签色"],
                [],
                ["二·B、吹气部分"],
                ["货名", "日产量/22H", "用料", "预估料重 g", "料价 HK$/lb", "产品料价", "吹工", "披锋", "小计", "利润 ×", "合计 HK$", "成品用量", "出数", "模价 (¥)", "备注"],
                ["吹气瓶", "12000", "ABS 750SW", 45, 8.5, 0.8425, 0.2, 0.1, 1.1425, 1.05, 1.1996, 1, "1出2", 5000, "透明"],
            ],
            title="啤机报价",
        ),
        "molding",
    )

    assert parsed.target_department == "molding"
    assert parsed.header_row == 2
    assert parsed.row_count == 2
    assert parsed.payload_fragment["injection_loss_rate_percent"] == "3.0000"
    assert parsed.payload_fragment["injection_lines"][0] == {
        "item": "主体模",
        "mold_no": "M-01",
        "material": "ABS",
        "grade": "750SW",
        "color": "黑色",
        "net_weight_g": "100.0000",
        "loss_rate_percent": "3.0000",
        "machine_name": "80T",
        "machine_code": "5A",
        "cavity": "2",
        "sets": "1.0000",
        "target_output": "5000.0000",
        "cycle_time_seconds": "24.0000",
        "quantity": "2.0000",
        "remark": "客签色",
        "source_row": 3,
    }
    assert parsed.payload_fragment["blow_lines"][0] == {
        "item": "吹气瓶",
        "daily_capacity": "12000",
        "material": "ABS",
        "grade": "750SW",
        "estimated_weight_g": "45.0000",
        "labor_hkd": "0.2000",
        "burr_hkd": "0.1000",
        "profit_multiplier": "1.0500",
        "quantity": "1.0000",
        "output_count": "1出2",
        "mold_price_rmb": "5000.0000",
        "remark": "透明",
        "source_row": 7,
    }
    assert any("不直接写入" in warning for warning in parsed.warnings)


def test_molding_import_replace_updates_both_sections_and_loss_rate():
    merged = _merge_import_payload(
        "molding",
        {
            "injection_loss_rate_percent": "5",
            "injection_lines": [{"item": "旧注塑"}],
            "blow_lines": [{"item": "旧吹气"}],
            "caixing_tool_plan_rows": [{"ref_no": "KEEP"}],
        },
        {
            "injection_loss_rate_percent": "3",
            "injection_lines": [{"item": "新注塑"}],
            "blow_lines": [{"item": "新吹气"}],
        },
        "replace",
    )

    assert merged["injection_loss_rate_percent"] == "3"
    assert [row["item"] for row in merged["injection_lines"]] == ["新注塑"]
    assert [row["item"] for row in merged["blow_lines"]] == ["新吹气"]
    assert merged["caixing_tool_plan_rows"] == [{"ref_no": "KEEP"}]


def test_painting_import_keeps_legacy_eight_operation_columns_and_row_metadata():
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                ["图片", "名称", "位置", "夹模", "夹模单价", "移印", "移印单价", "散枪", "散枪单价", "擦PP水", "擦PP水单价", "总报价", "备注"],
                ["头部.png", "公仔", "面部", 2, 0.5, 1, 0.3, 3, 0.2, 4, 0.1, 2.3, "对色板"],
            ]
        ),
        "painting",
    )

    row = parsed.payload_fragment["rows"][0]
    assert row["operations"]["clamp"] == {"quantity": "2.0000", "unit_price_hkd": "0.5000"}
    assert row["operations"]["pad_print"] == {"quantity": "1.0000", "unit_price_hkd": "0.3000"}
    assert row["operations"]["uv"] == {"quantity": "0.0000", "unit_price_hkd": "0.0000"}
    assert row["operations"]["spray"] == {"quantity": "3.0000", "unit_price_hkd": "0.2000"}
    assert row["operations"]["wipe"] == {"quantity": "0.0000", "unit_price_hkd": "0.0000"}
    assert row["operations"]["pp_water"] == {"quantity": "4.0000", "unit_price_hkd": "0.1000"}
    assert row["image_reference"] == "头部.png"
    assert row["name"] == "公仔"
    assert row["position"] == "面部"
    assert row["remark"] == "对色板"
    assert "总报价和合计仅用于核对" in parsed.warnings[0]


def test_painting_import_merges_a_two_row_operation_header():
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                ["位置", "夹模", None, "移印", None, "散枪", None, "备注"],
                [None, "数量", "单价", "数量", "单价", "数量", "单价", None],
                ["面部", 2, 0.5, 1, 0.3, 3, 0.2, "对色板"],
            ]
        ),
        "painting",
    )

    row = parsed.payload_fragment["rows"][0]
    assert parsed.header_row == 1
    assert parsed.row_count == 1
    assert row["source_row"] == 3
    assert row["operations"]["clamp"] == {"quantity": "2.0000", "unit_price_hkd": "0.5000"}
    assert row["operations"]["pad_print"] == {"quantity": "1.0000", "unit_price_hkd": "0.3000"}


def test_painting_download_template_maps_uv_and_shifted_operations_without_losing_formulas():
    content, _ = build_internal_quote_import_template("painting")
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active
    assert [sheet.cell(2, col).value for col in range(8, 12)] == ["UV", "UV单价", "散枪", "散枪单价"]
    assert "H3*I3" in sheet["V3"].value
    assert sheet["V41"].value == "=SUM(V3:V40)"
    assert len(sheet._images) == 14
    sheet["H3"], sheet["I3"] = 2, 0.52
    sheet["J3"], sheet["K3"] = 3, 0.08
    sheet["W3"] = "UV 与散枪分别核价"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    parsed = parse_internal_quote_workbook(stream.getvalue(), "painting")
    row = next(row for row in parsed.payload_fragment["rows"] if row["source_row"] == 3)
    assert row["operations"]["uv"] == {"quantity": "2.0000", "unit_price_hkd": "0.5200"}
    assert row["operations"]["spray"] == {"quantity": "3.0000", "unit_price_hkd": "0.0800"}
    assert row["remark"] == "UV 与散枪分别核价"
    from app.services.internal_quote_calculator import calculate_section
    result = calculate_section("painting", {"rows": [row]}, {}, "UV-TEST")
    assert result["totals"]["total_hkd"] == "1.5800"  # 移印 .30 + UV 1.04 + 散枪 .24


@pytest.mark.parametrize("two_row_header", [False, True])
def test_painting_import_accepts_uv_only_quotes_and_warns_about_missing_unit_price(two_row_header):
    headers = [["名称", "位置", "UV", "UV单价", "备注"]]
    if two_row_header:
        headers = [["名称", "位置", "UV", None, "备注"], [None, None, "数量", "单价", None]]
    parsed = parse_internal_quote_workbook(workbook_bytes(headers + [
        ["外壳", "正面", 2, .52, "UV"], ["外壳", "背面", 1, None, "待补价"],
    ]), "painting")
    assert parsed.row_count == 2
    assert parsed.payload_fragment["rows"][0]["operations"]["uv"]["unit_price_hkd"] == "0.5200"
    assert any("背面/UV 未识别单价" in warning for warning in parsed.warnings)


def test_slush_import_maps_visible_quote_fields_and_ignores_source_total():
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                ["产品编号", "胶件名称", "材料", "料重(g)", "日产量24H", "用量(PC)", "单价 HKD", "总价 HKD", "备注"],
                ["RC-01", "公仔手臂", "PVC", 35.5, 8000, 2, 3.6, 99, "透明"],
                ["合计", None, None, None, None, None, None, 99, None],
            ],
            title="搪胶报价",
        ),
        "slush",
    )

    assert parsed.target_department == "slush"
    assert parsed.header_row == 1
    assert parsed.row_count == 1
    assert parsed.payload_fragment["lines"][0] == {
        "product_code": "RC-01",
        "item": "公仔手臂",
        "material": "PVC",
        "weight_g": "35.5000",
        "daily_output_24h": "8000.0000",
        "quantity": "2.0000",
        "unit_price_hkd": "3.6000",
        "remark": "透明",
        "source_row": 2,
    }
    assert any("总价和合计仅用于核对" in warning for warning in parsed.warnings)


def test_slush_import_replace_only_replaces_slush_lines():
    merged = _merge_import_payload(
        "slush",
        {"lines": [{"product_code": "OLD"}], "legacy_note": "keep"},
        {"lines": [{"product_code": "RC-01"}]},
        "replace",
    )

    assert merged["lines"] == [{"product_code": "RC-01"}]
    assert merged["legacy_note"] == "keep"


def test_sewing_import_keeps_usage_rmb_price_markup_and_labor_line():
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                ["物料名称", "裁片部位", "供应商", "用量", "单价", "码点", "价钱", "备注"],
                ["绒布", "身体", "供应商A", 0.25, 12, 1.1, 3.3, "红色"],
                ["车缝人工", "", "", 1, 2, 1, 2, ""],
            ]
        ),
        "sewing",
    )

    group = parsed.payload_fragment["groups"][0]
    assert parsed.target_department == "sewing"
    assert group["category"] == "clothes"
    assert group["materials"][0]["usage"] == "0.2500"
    assert group["materials"][0]["unit_price_rmb"] == "12.0000"
    assert group["materials"][0]["markup"] == "1.1000"
    assert group["materials"][1]["item"] == "车缝人工"
    assert group["labor_rmb"] == "0.0000"


def test_sewing_import_maps_reference_template_exchange_rate_and_record_fields():
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                ["物料名称", "裁片部位", "供应商", "布料MOQ/Y", "低于MOQ/每色产生费用", "用量/码", "单价", "汇率", "成本", "码点", "价钱", "备注"],
                ["8寸紫色土豆蝙蝠"],
                ['58"270G白色莱卡布', "前身", "恒欣", 500, 20, 0.084, 93.2, 0.85, 9.2075, 1.1, 10.1282, "感温变色"],
            ],
            title="明细260707",
        ),
        "sewing",
    )

    row = parsed.payload_fragment["groups"][0]["materials"][0]
    assert parsed.sheet_name == "明细260707"
    assert parsed.payload_fragment["groups"][0]["name"] == "8寸紫色土豆蝙蝠"
    assert row["supplier"] == "恒欣"
    assert row["fabric_moq_y"] == "500.0000"
    assert row["below_moq_fee_rmb"] == "20.0000"
    assert row["usage"] == "0.0840"
    assert row["unit_price_rmb"] == "93.2000"
    assert row["exchange_rate"] == "0.8500"
    assert row["markup"] == "1.1000"
    assert any("单价 RMB ÷ 行汇率" in warning for warning in parsed.warnings)


def test_sewing_import_forward_fills_material_for_following_cutting_parts():
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                ["物料名称", "裁片部位", "供应商", "布料MOQ/Y", "低于MOQ/每色产生费用", "用量/码", "单价", "成本", "码点", "价钱", "备注"],
                ["8寸紫色土豆蝙蝠"],
                ['58\"270G白色莱卡布，感温变色涂层', "前身", "恒欣", 500, None, 0.084, 93.2, 7.87, 1.1, 8.66, ""],
                [None, "后身", None, None, None, 0.085, 93.2, 7.97, 1.1, 8.76, ""],
                [None, "底部", None, None, None, 0.034, 93.2, 3.2, 1.1, 3.52, ""],
                ['58\"270G紫色莱卡布(无温变)', "前朵", "恒欣", 50, None, 0.008, 20.7, 0.16, 1.1, 0.18, ""],
                [None, "后耳", None, None, None, 0.008, 20.7, 0.16, 1.1, 0.18, ""],
            ],
            title="明细",
        ),
        "sewing",
    )

    materials = parsed.payload_fragment["groups"][0]["materials"]
    assert parsed.row_count == 5
    assert parsed.payload_fragment["groups"][0]["name"] == "8寸紫色土豆蝙蝠"
    assert [row["item"] for row in materials] == [
        '58\"270G白色莱卡布，感温变色涂层',
        '58\"270G白色莱卡布，感温变色涂层',
        '58\"270G白色莱卡布，感温变色涂层',
        '58\"270G紫色莱卡布(无温变)',
        '58\"270G紫色莱卡布(无温变)',
    ]
    assert [row["part"] for row in materials] == ["前身", "后身", "底部", "前朵", "后耳"]
    assert materials[1]["usage"] == "0.0850"
    assert materials[1]["unit_price_rmb"] == "93.2000"
    assert materials[1]["markup"] == "1.1000"


def test_sewing_import_preserves_source_precision_and_uses_total_row_as_product_boundary():
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                ["物料名称", "裁片部位", "供应商", "布料MOQ/Y", "低于MOQ/每色产生费用", "用量/码", "单价", "成本", "码点", "价钱", "备注"],
                ["8寸紫色土豆蝙蝠"],
                ['58"270G白色莱卡布，感温变色涂层', "前身", "恒欣", 500, None, 0.110590277777778, 93.1, 10.2959548611111, 1.1, 11.3255503472222, ""],
                [None, "后身", None, None, None, 0.116666666666667, 93.1, 10.8616666666667, 1.1, 11.9478333333333, ""],
                ["25 mm黄色纽扣（兴信提供）"],
                [None, None, None, None, None, None, None, None, "合计", 23.2733836805555, None],
                ["10寸紫色小狗"],
                ["紫色莱卡布", "前朵", "恒欣", 50, None, 0.01, 20.7, 0.207, 1.1, 0.2277, ""],
            ],
            title="明细",
        ),
        "sewing",
    )

    assert parsed.row_count == 4
    assert [group["name"] for group in parsed.payload_fragment["groups"]] == ["8寸紫色土豆蝙蝠", "10寸紫色小狗"]
    first_group = parsed.payload_fragment["groups"][0]["materials"]
    assert [row["item"] for row in first_group] == [
        '58"270G白色莱卡布，感温变色涂层',
        '58"270G白色莱卡布，感温变色涂层',
        "25 mm黄色纽扣（兴信提供）",
    ]
    assert [row["part"] for row in first_group[:2]] == ["前身", "后身"]
    assert first_group[0]["usage"] == "0.110590277777778"
    assert first_group[1]["usage"] == "0.116666666666667"
    assert first_group[2]["usage"] == "1.0000"
    assert first_group[2]["unit_price_rmb"] == "0.0000"


def test_sewing_import_maps_screenshot_fields_multiple_groups_and_ignores_derived_totals():
    header = ["布料名称", "部位", "工艺", "裁片数", "用量/码", "物料价(RMB)", "价钱(RMB)", "码点", "总价钱(RMB)", "备注"]
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                ["产品：6寸小蜥蜴 / 盾牌"],
                header,
                ["绒布", "身体", "电绣", 4, 0.25, 12, 3, 1.1, 99, "红色"],
                ["车缝人工", "", "", 0, 1, 2, 2, 1, 99, ""],
                ["合计", None, None, None, None, None, None, None, 198, None],
                ["产品：发套"],
                header,
                ["毛料", "头部", "", 2, 0.5, 8, 4, 0, 88, "黑色"],
            ],
            title="车缝报价",
        ),
        "sewing",
    )

    assert parsed.target_department == "sewing"
    assert parsed.header_row == 2
    assert parsed.row_count == 3
    assert [group["name"] for group in parsed.payload_fragment["groups"]] == ["6寸小蜥蜴 / 盾牌", "发套"]
    first = parsed.payload_fragment["groups"][0]["materials"][0]
    assert first == {
        "item": "绒布",
        "part": "身体",
        "craft": "电绣",
        "pieces": "4.0000",
        "supplier": "",
        "usage": "0.2500",
        "unit_price_rmb": "12.0000",
        "markup": "1.1000",
        "remark": "红色",
        "source_row": 3,
    }
    assert parsed.payload_fragment["groups"][1]["category"] == "hair"
    assert parsed.payload_fragment["groups"][1]["materials"][0]["markup"] == "1.0000"
    assert any("总价钱和合计仅用于核对" in warning for warning in parsed.warnings)


def test_sewing_import_preserves_screen_printing_craft():
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
                [
                    ["布料名称", "部位", "工艺", "裁片数", "用量/码", "物料价(RMB)", "价钱(RMB)", "码点"],
                    ["网布", "正面", "丝印", 1, 0.5, 6, 3, 1],
            ],
            title="车缝报价",
        ),
        "sewing",
    )

    material = parsed.payload_fragment["groups"][0]["materials"][0]
    assert material["craft"] == "丝印"


def test_assembly_import_builds_process_groups_and_uses_quote_qty_fallback():
    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                ["组装工序名称", "人数", "工具", "备注"],
                ["装电池", 4, "电批", "检查极性"],
            ]
        ),
        "assembly",
        fallback_qty=Decimal("1000"),
    )

    group = parsed.payload_fragment["groups"][0]
    process = group["processes"][0]
    assert group["category"] == "assembly"
    assert group["production_qty"] == "1000.0000"
    assert group["teams"] == "1.0000"
    assert process["name"] == "装电池"
    assert process["persons"] == "4.0000"
    assert process["remark"] == "检查极性"
    assert any("报价数量" in warning for warning in parsed.warnings)


def test_assembly_import_maps_workshop_regions_and_switches_to_packaging_from_column_c():
    header = [
        "报客价/港币",
        "报价人",
        "货号或图片",
        "做工名称",
        "总目标数量",
        "人数",
        "写报表的工价(￥)",
        "正班时间",
        "正班生产目标数",
        "正班8小时工资",
        "加班时间",
        "加班生产数",
        "加班工资",
        "合计工资(含加班)",
        "用量",
        "合计工价($)",
        "报客工价(HK$)",
        "报价人工标准",
        "车间标准工资(不含加班)",
        "报价日期",
        "备注（看图、手办、样板、参考办）",
    ]

    def detail(
        group: str | None,
        process: str | None,
        production_qty: int | None,
        persons: int | None,
        remark: str = "手办报价",
    ) -> list[object]:
        row: list[object] = [None] * len(header)
        row[2] = group
        row[3] = process
        row[4] = production_qty
        row[5] = persons
        row[20] = remark
        return row

    parsed = parse_internal_quote_workbook(
        workbook_bytes(
            [
                header,
                ["车间填写", "车间填写", "车间填写", "车间填写", "车间填写", "车间填写"],
                detail("组装马桶", "测试IC板", 3000, 1),
                detail(None, "焊喇叭", 3000, 2),
                detail(None, None, None, 3, ""),
                detail("组装马桶", "剪软管", 3000, 1),
                detail(None, "装箱/杂工", 3000, 2),
                detail(None, None, None, 3, ""),
                detail("包装公仔", "彩盒印日期码", 3000, 1),
                detail(None, "折彩盒", 3000, 6),
                detail(None, None, None, 7, ""),
                detail("出口配件", "封箱", None, 2),
            ],
            title="组装",
        ),
        "assembly",
        fallback_qty=Decimal("5000"),
    )

    groups = parsed.payload_fragment["groups"]
    assert parsed.header_row == 1
    assert parsed.row_count == 7
    assert [group["name"] for group in groups] == [
        "组装马桶",
        "组装马桶",
        "包装公仔",
        "出口配件",
    ]
    assert [group["category"] for group in groups] == [
        "assembly",
        "assembly",
        "packaging",
        "packaging",
    ]
    assert [group["production_qty"] for group in groups] == [
        "3000.0000",
        "3000.0000",
        "3000.0000",
        "5000.0000",
    ]
    assert [row["persons"] for row in groups[0]["processes"]] == ["1.0000", "2.0000"]
    assert groups[1]["processes"][1]["name"] == "装箱/杂工"
    assert groups[1]["processes"][1]["remark"] == "手办报价"
    assert any("出口配件" in warning and "报价数量" in warning for warning in parsed.warnings)
    assert any("C 列区域" in warning and "F 列工艺人数" in warning for warning in parsed.warnings)


def test_binary_xls_is_rejected_by_p3_parser():
    with pytest.raises(ValueError, match="xlsx/xlsm"):
        parse_internal_quote_workbook(b"\xD0\xCF\x11\xE0fake", "mold")
