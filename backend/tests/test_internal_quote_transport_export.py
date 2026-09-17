import json
import re
from decimal import Decimal, ROUND_HALF_UP
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.services.internal_quote import _rr2_cost_summary
from app.services.internal_quote_calculator import calculate_section
from app.services.internal_quote_excel import _build_summary_sheet, _purchase_detail_label, _write_transport_inputs
from test_internal_quote_excel_formula_links import _section, _find_row


def cell_value(sheet, address):
    value = sheet[address.replace('$', '')].value
    if not isinstance(value, str) or not value.startswith('='):
        return value or 0
    expression = re.sub(r'\$?([A-Z]+)\$?(\d+)', lambda m: repr(cell_value(sheet, m[1] + m[2])), value[1:])
    funcs = {'IF': lambda condition, yes, no: yes if condition else no,
             'MAX': max, 'ROUND': lambda x, places: float(Decimal(str(x)).quantize(Decimal(10) ** -places, rounding=ROUND_HALF_UP)),
             'VALUE': float, 'LEFT': lambda s, n: str(s)[:n], 'RIGHT': lambda s, n: str(s)[-n:],
             'LEN': lambda s: len(str(s)), 'FIND': lambda part, s: str(s).index(part) + 1}
    return eval(expression, {'__builtins__': {}}, funcs)


@pytest.mark.parametrize('mode', ['ordinary', 'component'])
@pytest.mark.parametrize('unit', ['inch', 'cm'])
@pytest.mark.parametrize('inner', [False, True])
@pytest.mark.parametrize('fees', ['freight', 'lifting', 'both'])
def test_freight_and_cbm_follow_visible_carton_inputs(mode, unit, inner, fees):
    snapshot = {'fx': {'rmb_hkd': '.85', 'hkd_usd': '7.8'},
                'freight': {'routes': [
                    {'route_key': 'hk40', 'route_name': 'HK 40 柜', 'capacity_key': 'cap_40', 'freight_hkd': '4600', 'lifting_hkd': '300'},
                    {'route_key': 'hk20', 'route_name': 'HK 20 柜', 'capacity_key': 'cap_20', 'freight_hkd': '4300', 'lifting_hkd': '200'}]}}
    payload = {'cartons': [{'item': '外箱', 'length_in': 25.75, 'width_in': 13.75, 'height_in': 15, 'size_unit': unit, 'qty_per_carton': 6}],
               'paper_price_factor': 2.75, 'testing_fee_enabled': False,
               'freight_calc': {'enabled': True, 'freight_enabled': fees != 'lifting', 'lifting_enabled': fees != 'freight',
                                'selected_route_keys': ['hk40'], 'cap_40': 1980},
               'shipping': {'markup_x': 1.2, 'misc_ratio': .03}}
    if inner:
        payload['cartons'].append({'item': '内箱', 'length_in': 10, 'width_in': 5, 'height_in': 5, 'qty_per_carton': 2, 'size_unit': unit})
    if mode == 'component':
        payload.update(pricing_mode='component', pricing_components=[{'id': 'a', 'name': '主体', 'markup_x': 1.2}],
                       color_box_size_in={'length': 25, 'width': 13, 'height': 14})
    calculation = calculate_section('sales', payload, snapshot, 'TRANSPORT')
    assert calculation['status'] == 'valid', calculation.get('warnings')
    section = _section('sales', payload=payload, calculation=calculation)
    section.is_required = True
    context = {'carton_hkd': Decimal(calculation['totals']['carton_hkd']), 'factory_price_hkd': Decimal(calculation['totals']['carton_hkd'])}
    summary = _rr2_cost_summary([section], context, snapshot, factory_id='huakang-b' if mode == 'component' else 'huaxing')
    workbook = Workbook()
    _build_summary_sheet(workbook, SimpleNamespace(product_name='运输公式验证', quote_no='TRANSPORT', customer='JustPlay' if mode == 'component' else '普通客', region_code='mainland', remark=''),
                         [section], snapshot, summary, context, [])
    sheet = workbook['报价明细']
    cbm = next(c for row in sheet for c in row if c.value == 'CBM:')
    cuft = sheet.cell(cbm.row - 1, cbm.column + 1)
    cbm_value = sheet.cell(cbm.row, cbm.column + 1)
    assert cell_value(sheet, cbm_value.coordinate) == pytest.approx(25.75 * 13.75 * 15 * .0254**3)
    assert '!' not in cbm_value.value and cuft.coordinate in cbm_value.value
    cartons_row = _find_row(sheet, 3, '装柜箱数')
    assert cell_value(sheet, f'E{cartons_row}') == 644
    expected = calculation['totals']['freight_options'][0]
    fee_cells = []
    for enabled, label, key in [(fees != 'lifting', '运费', 'freight_per_piece_hkd'), (fees != 'freight', '吊柜费', 'lifting_per_piece_hkd')]:
        if enabled:
            row = _find_row(sheet, 3, label)
            cell = sheet.cell(row, 5)
            assert cell.data_type == 'f' and '!' not in cell.value
            assert cell_value(sheet, cell.coordinate) == float(expected[key])
            fee_cells.append((cell, float(expected[key])))
    assert not any(c.value == 'HK 20 柜' for row in sheet for c in row)
    packing_row = next(c.row for row in sheet for c in row if c.value == '装箱：')
    sheet.cell(packing_row, cbm.column + 1).value = '2/12' if inner else 12
    for cell, previous in fee_cells:
        assert cell_value(sheet, cell.coordinate) == pytest.approx(previous / 2, abs=.0001)
    # Edit the dimensions/source that owns the visible outer carton.
    dimension_label = '彩盒尺寸 (inch)' if mode == 'component' else f'外箱 ({unit}):'
    dimension_row = next(c.row for row in sheet for c in row if c.value == dimension_label)
    dimension = sheet.cell(dimension_row, cbm.column + 1)
    dimension.value = float(dimension.value) + 10
    assert cell_value(sheet, cbm_value.coordinate) > 25.75 * 13.75 * 15 * .0254**3
    assert cell_value(sheet, f'E{cartons_row}') < 644
    workbook.close()


@pytest.mark.parametrize('entry,expected', [
    ({'kind': 'material', 'label': '螺丝', 'specification': '2.6*8', 'quantity': 3}, '2.6*8 螺丝（3pcs）'),
    ({'kind': 'material', 'label': '2.6*8 螺丝', 'specification': '2.6*8', 'quantity': 1}, '2.6*8 螺丝（1pcs）'),
    ({'kind': 'packaging_material', 'label': '绳子', 'quantity': .25, 'unit': 'm'}, '绳子（0.25m）'),
    ({'kind': 'material', 'label': '旧配件'}, '旧配件'),
])
def test_material_description_preserves_specification_and_real_quantity(entry, expected):
    assert _purchase_detail_label(entry) == expected


@pytest.mark.parametrize('freight_enabled,lifting_enabled', [(True, False), (False, True), (True, True)])
def test_legacy_shared_transport_total_keeps_frozen_fee_allocation(freight_enabled, lifting_enabled):
    # Old freight options split the saved per-piece total by 48%/52%.
    # Their whole-container price cannot be treated as the freight-only input.
    options = [{'item': '港柜', 'capacity_cuft': '1980', 'freight_cost_hkd': '4600',
                'lifting_cost_hkd': '0', 'has_lifting_fee': False, 'per_piece_hkd': '1.1905'}]
    routes = [{'name': '港柜', 'freight_hkd': .5714 if freight_enabled else 0,
               'lift_hkd': .6191 if lifting_enabled else 0}]
    section = _section('sales', calculation={'totals': {'freight_options': options}})
    workbook = Workbook()
    next_row, values = _write_transport_inputs(workbook.active, 10, routes, section, '$O$31', 'O35',
                                               freight_enabled=freight_enabled, lifting_enabled=lifting_enabled)
    assert next_row == 10
    assert values == [{'freight_hkd': routes[0]['freight_hkd'], 'lift_hkd': routes[0]['lift_hkd']}]
    assert workbook.active['E10'].value is None
    workbook.close()
