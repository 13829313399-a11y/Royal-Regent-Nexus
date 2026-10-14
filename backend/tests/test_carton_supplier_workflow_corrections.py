"""Current-version supplier handoff and non-arrival corrections, in isolated databases."""
import json
from io import BytesIO
from uuid import uuid4
from test_molding_sample_api import make_client, login_as
from test_carton_procurement_api import _order_payload, _submit_order
from test_carton_supplier_portal import (BASE, setup_portal, supplier_login, accept_all,
    ship_payload, receive_payload, revoke_supplier_permission, restore_supplier_permission)


def workspace_order(client, identifier=None):
    orders = client.get(BASE + "/workspace", params={"factory_id": "huaxing"}).json()["orders"]
    return next(order for order in orders if identifier is None or order["id"] == identifier)


def test_supplier_promise_defer_and_non_arrival_reopen_preserve_single_receipt(monkeypatch):
    monkeypatch.setenv("AUTHZ_MODE", "legacy")
    with make_client(monkeypatch) as client:
        setup_portal(client)
        order = accept_all(client)
        current = workspace_order(client)
        paper = current["lines"][0]
        promise = {"factory_id": "huaxing", "issue_id": current["issue_id"],
                   "expected_revision": paper["commitment_revision"], "promised_date": "2026-09-28"}
        changed = client.put(BASE + f"/papers/{paper['id']}/commitment", json=promise)
        assert changed.status_code == 200, changed.text
        assert changed.json()["planned_date"] == current["planned_date"]
        assert changed.json()["lines"][0]["promised_date"] == "2026-09-28"
        assert client.put(BASE + f"/papers/{paper['id']}/commitment", json=promise).status_code == 409
        shipment = client.post(BASE + "/shipments", json=ship_payload(order)).json()
        login_as(client, "warehouse_keeper")
        action = {"factory_id": "huaxing", "expected_revision": shipment["revision"],
                  "request_id": str(uuid4()), "reason": "车辆尚未到达"}
        from app.services.auth import AuthContext, AuthGrantContext, AuthOverrideContext, get_current_user
        permissions = frozenset({"carton_procurement:read", "carton_procurement:receipt_write", "carton_procurement:inventory_write"})
        denied_actor = AuthContext("user-warehouse", "warehouse_keeper", "仓管", (), (), permissions,
            ("huaxing",), ("pmc-warehouse",), grants=(AuthGrantContext("test-role", "仓管", "huaxing", "pmc-warehouse", permissions),),
            overrides=(AuthOverrideContext("deny-inventory", "carton_procurement:inventory_write", "deny", "*", "*"),))
        # Canonical deny applies even when legacy mode would grant the old position permissions.
        client.app.dependency_overrides[get_current_user] = lambda: denied_actor
        assert client.post(BASE + f"/internal/shipments/{shipment['id']}/defer", json=action).status_code == 403
        from app.db import SessionLocal
        from app.models.carton_supplier_portal import SupplierShipment
        with SessionLocal() as db:
            stored = db.get(SupplierShipment, shipment["id"])
            assert stored.status == "SENT" and stored.revision == shipment["revision"]
        client.app.dependency_overrides.clear()
        deferred = client.post(BASE + f"/internal/shipments/{shipment['id']}/defer", json=action)
        assert deferred.status_code == 200, deferred.text
        assert deferred.json()["status"] == "SENT"
        assert deferred.json()["pending_reason"] == "车辆尚未到达"
        assert client.post(BASE + f"/internal/shipments/{shipment['id']}/defer", json=action).json()["revision"] == deferred.json()["revision"]
        zero = receive_payload(client, deferred.json())
        for line in zero["lines"]:
            line.update(received_quantity=0, location_allocations=[], difference_reason="误确认全部未到")
        absent = client.post(BASE + f"/internal/shipments/{shipment['id']}/receive", json=zero)
        assert absent.status_code == 200, absent.text
        assert absent.json()["status"] == "NOT_RECEIVED" and not absent.json()["receipt_id"]
        reopen = dict(action, expected_revision=absent.json()["revision"], request_id=str(uuid4()), reason="原未到货反馈有误")
        client.app.dependency_overrides[get_current_user] = lambda: denied_actor
        assert client.post(BASE + f"/internal/shipments/{shipment['id']}/reopen", json=reopen).status_code == 403
        with SessionLocal() as db:
            stored = db.get(SupplierShipment, shipment["id"])
            assert stored.status == "NOT_RECEIVED" and stored.revision == absent.json()["revision"]
        client.app.dependency_overrides.clear()
        reopened = client.post(BASE + f"/internal/shipments/{shipment['id']}/reopen", json=reopen)
        assert reopened.status_code == 200, reopened.text
        assert reopened.json()["status"] == "SENT" and reopened.json()["acceptance_lines"] == []
        assert any(row["status"] == "NOT_RECEIVED" for row in reopened.json()["acceptance_history"])
        # Old all-zero request remains an unchanged replay; it cannot terminalize the reopened note.
        old = client.post(BASE + f"/internal/shipments/{shipment['id']}/receive", json=zero)
        assert old.status_code == 200 and old.json()["status"] == "SENT", old.text
        actual = receive_payload(client, reopened.json())
        posted = client.post(BASE + f"/internal/shipments/{shipment['id']}/receive", json=actual)
        assert posted.status_code == 200 and posted.json()["receipt_id"], posted.text
        replay = client.post(BASE + f"/internal/shipments/{shipment['id']}/receive", json=actual)
        assert replay.status_code == 200 and replay.json()["receipt_id"] == posted.json()["receipt_id"]
        assert client.post(BASE + f"/internal/shipments/{shipment['id']}/reopen", json=dict(reopen, request_id=str(uuid4()), expected_revision=posted.json()["revision"])).status_code == 409
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonAuditEvent, CartonReceipt
        from sqlalchemy import select
        with SessionLocal() as db:
            receipts = db.scalars(select(CartonReceipt)).all()
            assert len(receipts) == 1
            events = db.scalars(select(CartonAuditEvent).where(CartonAuditEvent.event_type == "SUPPLIER_PAPER_ACCEPTED")).all()
            assert any(json.loads(event.detail_json).get("previous_promised_date") == "2026-09-23" for event in events)


def test_cancellation_requires_current_supplier_ack_and_approval(monkeypatch):
    with make_client(monkeypatch) as client:
        original = setup_portal(client)
        accept_all(client)
        login_as(client, "admin")
        cancelled = client.post(f"/api/carton-procurement/orders/{original['order_no']}/reduce", json={
            "factory_id": "huaxing", "expected_revision": original["revision"], "reduction_quantity": 3600,
            "reason": "客户取消整张订单"})
        assert cancelled.status_code == 200, cancelled.text
        supplier_login(client)
        order = workspace_order(client)
        assert order["status"] == "CANCELLED"
        notice = order["cancellation_notice"]
        assert notice["issued"] and not notice["acknowledged"]
        assert all(float(line["required_quantity"]) == 0 for line in order["lines"])
        documents = client.get(BASE + "/documents", params={"factory_id": "huaxing"}).json()
        cancellation = next(document for document in documents if document["id"] == notice["issue_id"])
        assert cancellation["is_cancellation"]
        erp = client.post(BASE + "/documents/order-import.xlsx", json={"documents": [{"factory_id": "huaxing", "kind": "PURCHASE", "id": notice["issue_id"]}], "acknowledge_unaccepted": True})
        assert erp.status_code == 422 and "取消采购通知" in erp.text, erp.text
        exported = client.post(BASE + "/documents/export.xlsx", json={"documents": [{"factory_id": "huaxing", "kind": "PURCHASE", "id": notice["issue_id"]}]})
        assert exported.status_code == 200, exported.text
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(exported.content)); assert "取消采购通知" in book.active["D2"].value
        payload = {"factory_id": "huaxing", "issue_id": notice["issue_id"], "expected_order_revision": notice["order_revision"]}
        assert client.post(BASE + f"/orders/{order['id']}/cancellation-ack", json=dict(payload, expected_order_revision=notice["order_revision"] - 1)).status_code == 409
        revoke_supplier_permission("carton_supplier:approve")
        supplier_login(client)
        assert client.post(BASE + f"/orders/{order['id']}/cancellation-ack", json=payload).status_code == 403
        restore_supplier_permission("carton_supplier:approve")
        supplier_login(client)
        accepted = client.post(BASE + f"/orders/{order['id']}/cancellation-ack", json=payload)
        assert accepted.status_code == 200 and accepted.json()["cancellation_notice"]["acknowledged"], accepted.text
        assert client.post(BASE + f"/orders/{order['id']}/cancellation-ack", json=payload).status_code == 200
        assert client.post(BASE + "/shipments", json=ship_payload(order)).status_code in {404, 409}


def test_ambiguous_delivery_selection_rechecks_identity_revision_and_replay(monkeypatch):
    from openpyxl import Workbook
    with make_client(monkeypatch) as client:
        setup_portal(client)
        first = accept_all(client)
        login_as(client, "admin")
        body = _order_payload()
        body.update(customer_po="SECOND-PO", request_id=str(uuid4()))
        second = client.post("/api/carton-procurement/orders", json=body)
        assert second.status_code == 201, second.text
        second = _submit_order(client, second.json())
        supplier_login(client)
        current = workspace_order(client, second["id"])
        for line in current["lines"]:
            accepted = client.put(BASE + f"/papers/{line['id']}/commitment", json={"factory_id": "huaxing",
                "issue_id": current["issue_id"], "expected_revision": 0, "promised_date": "2026-09-25"})
            assert accepted.status_code == 200, accepted.text
        current = workspace_order(client, second["id"])
        paper = current["lines"][0]
        book = Workbook(); sheet = book.active; sheet.title = "送货明细"
        sheet.append(["送货单号", "日期", "客户", "客户单号", "客户料号", "名称", "纸质", "规格", "数量"])
        sheet.append(["DN-CHOOSE", "2026-09-24", "华兴", current["contract_no"], current["item_no"],
                      paper["packaging_type"], paper["paper_quality"], paper["specification"], 5])
        output = BytesIO(); book.save(output); upload = {"file": ("送货.xlsx", output.getvalue())}
        for malformed in ({}, "", False, 0):
            rejected = client.post(BASE + "/shipments/import-preview", files=upload, data={"row_selections": json.dumps(malformed)})
            assert rejected.status_code == 422, rejected.text
        preview = client.post(BASE + "/shipments/import-preview", files=upload)
        assert preview.status_code == 200, preview.text
        row = preview.json()["groups"][0]["rows"][0]
        assert row["status"] == "BLOCKED" and len(row["candidates"]) == 2
        choice = [{"source_sheet": row["source_sheet"], "source_row": row["source_row"],
                   "order_line_id": paper["id"], "issue_id": current["issue_id"], "order_revision": current["revision"]}]
        reviewed = client.post(BASE + "/shipments/import-preview", files=upload, data={"row_selections": json.dumps(choice)})
        assert reviewed.status_code == 200 and reviewed.json()["groups"][0]["ready"], reviewed.text
        form = {"sha256": preview.json()["sha256"], "selections": json.dumps([{"factory_id": "huaxing", "delivery_note_no": "DN-CHOOSE"}]), "row_selections": json.dumps(choice)}
        stale = [dict(choice[0], order_revision=current["revision"] + 1)]
        assert client.post(BASE + "/shipments/import-confirm", files=upload, data=dict(form, row_selections=json.dumps(stale))).status_code == 409
        foreign = [dict(choice[0], order_line_id="OTHER-FACTORY-LINE")]
        assert client.post(BASE + "/shipments/import-confirm", files=upload, data=dict(form, row_selections=json.dumps(foreign))).status_code == 409
        confirmed = client.post(BASE + "/shipments/import-confirm", files=upload, data=form)
        assert confirmed.status_code == 200, confirmed.text
        shipment = confirmed.json()["shipments"][0]
        assert shipment["lines"][0]["order_line_id"] == paper["id"]
        assert client.post(BASE + "/shipments/import-confirm", files=upload, data=form).json()["shipments"][0]["id"] == shipment["id"]
        assert client.post(BASE + "/shipments/import-confirm", files=upload, data=dict(form, row_selections="[]")).status_code == 409
