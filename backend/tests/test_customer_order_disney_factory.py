from copy import deepcopy
from io import BytesIO

import openpyxl
import pytest
from fastapi import HTTPException
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from app.api import customer_order as api
from app.services import customer_order_unified as service
from test_customer_order_huaxing_unified import template, data


def disney_po():
    text = '''ORDER NUMBER ORIGINAL / CONFIRMATION PAGE
L-9743347 CHANGE 1
EARLIEST SHIP DATE LATEST SHIP DATE ORDER DATE SUPPLIER NUMBER
11/30/2026 12/7/2026 5/21/2026 40059618
REVISION NUMBER
1000128076 MM ASTRONAUT PULLBACK 2502 EA 1/EA 2.1400 401060937150 14.99
Multi INNER/PK: 6/EA
No Size CASE/PK: 6'''
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                             NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    stream.set_data(('BT /F1 9 Tf 12 TL 30 760 Td ' + ' T* '.join(f'({line}) Tj' for line in text.splitlines()) + ' ET').encode())
    page[NameObject('/Contents')] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return [('L-9743347.pdf', output.getvalue())]


def arguments(factory, po_files):
    wb = template('disney')
    ws = wb['ITEM表']
    ws['E4'], ws['H4'], ws['I4'] = factory + '-OLD', '1000128076', factory + '-产品'
    ws['K4'], ws['L4'] = 120, 6
    return dict(customer_code='disney', factory_id=factory, received_date='2026-09-08',
                po_files=po_files, schedule_file_name=factory + '.xlsx', schedule_content=data(wb))


@pytest.mark.parametrize('factory', ['huaxing', 'huakang-d'])
def test_disney_same_parser_exports_only_selected_factory_workbook(factory):
    api._ensure_customer_factory('disney', factory)
    assert api._get_mapped_customer_spec('disney', factory).po_extensions == ('.pdf',)
    args = arguments(factory, disney_po())
    preview = service.create_unified_customer_preview(**args)
    row = preview['rows'][0]
    assert preview['factory_id'] == factory
    assert row['product_name_zh'] == factory + '-产品'
    assert (row['contract_no'], row['reference_no'], row['po_no']) == ('', 'L-9743347', 'L-9743347')
    assert row['target_template'] == preview['target_template'] == service.huaxing.target_template(factory)
    assert ('华康D' if factory == 'huakang-d' else '华兴') in row['lineage']['mapping_rule']
    keys = {i['skip_key'] for r in preview['rows'] for i in r['issues'] if i.get('skip_key')}
    output, _, exported = service.export_unified_customer_schedule(**args, skipped_issue_keys=keys)
    assert exported['factory_id'] == factory
    wb = openpyxl.load_workbook(BytesIO(output))
    ws = wb['ITEM表']
    assert ws['E4'].value == factory + '-OLD'
    assert ws['E5'].value == ws['F5'].value == 'L-9743347'
    assert ws['I5'].value == factory + '-产品'
    assert ws['K5'].value == 2502
    for name in ['接单表', '正单评审表']:
        assert wb[name]['D4'].value == 'L-9743347'
        assert "'ITEM表'!K5" in wb[name]['J4'].value
    other_factory = 'huaxing' if factory == 'huakang-d' else 'huakang-d'
    assert all(other_factory not in str(c.value) for sheet in wb for rows in sheet for c in rows)
    assert openpyxl.load_workbook(BytesIO(args['schedule_content']))['ITEM表']['E5'].value is None


@pytest.mark.parametrize('factory', ['huakang-a', 'huakang-c', 'huadeng'])
def test_disney_unregistered_factories_remain_rejected(factory):
    with pytest.raises(HTTPException) as error:
        api._ensure_customer_factory('disney', factory)
    assert error.value.status_code == 400
    with pytest.raises(service.CustomerOrderUnifiedError, match='不属于当前厂区'):
        service.create_unified_customer_preview(**arguments(factory, disney_po()))


def test_disney_cross_factory_fingerprint_cannot_authorize_export():
    args = arguments('huaxing', disney_po())
    preview = service.create_unified_customer_preview(**args)
    fingerprint = api._preview_fingerprint(preview, args['received_date'])
    # Even identical files and mapping identities cannot reuse a different factory's approval.
    changed = deepcopy(preview)
    changed['factory_id'] = 'huakang-d'
    with pytest.raises(HTTPException) as error:
        api._ensure_preview_fingerprint(changed, args['received_date'], fingerprint)
    assert error.value.status_code == 409


def test_disney_other_factory_history_does_not_mark_duplicate():
    po_files = disney_po()
    hx_args = arguments('huaxing', po_files)
    wb = openpyxl.load_workbook(BytesIO(hx_args['schedule_content']))
    wb['ITEM表']['E4'], wb['ITEM表']['K4'] = 'L-9743347', 2502
    hx_args['schedule_content'] = data(wb)
    hx = service.create_unified_customer_preview(**hx_args)
    hd = service.create_unified_customer_preview(**arguments('huakang-d', po_files))
    assert any(i['code'] == 'duplicate_existing_order' for i in hx['rows'][0]['issues'])
    assert not any(i['code'] == 'duplicate_existing_order' for i in hd['rows'][0]['issues'])
