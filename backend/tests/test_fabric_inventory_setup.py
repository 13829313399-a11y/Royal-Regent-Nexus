import io
import importlib
from datetime import date
from decimal import Decimal
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from test_molding_sample_api import login_as
from warehouse_location_fixtures import make_client
from test_warehouse_operations import payload, post, workspace, quality

BASE = '/api/warehouse-operations/fabric/inventory-setup'
HEADERS = ['来源行号','物料编码','物料名称','物料分类','单位','仓库','仓位','缸号','卷号','期初数量','来源单位','备注']


def book(rows=None):
    workbook = Workbook(); sheet = workbook.active; sheet.title='期初库存'; sheet.append(HEADERS)
    for row in rows or [line()]:
        sheet.append(row)
    output=io.BytesIO();workbook.save(output);return output.getvalue()


def line(**changes):
    row=dict(zip(HEADERS,['001','00012','白色布','布料','码','布料一仓','A01','00001','0001','10.000001','旧供应商','上线核实']))
    row.update(changes);return list(row.values())


def preview(client, content, **changes):
    return client.post(BASE+'/preview',data=dict(factory_id='huakang-c',cutoff='2026-07-01',source_no='旧账001',**changes),files={'file':('opening.xlsx',content)})


def apply(client, content, result, request_id=None, **changes):
    params=dict(factory_id='huakang-c',cutoff='2026-07-01',source_no='旧账001',request_id=request_id or str(uuid4()),token=result['token'],confirmed='true')
    params.update(changes)
    return client.post(BASE+'/import',data=params,files={'file':('renamed.xlsx',content)})


def count(client, ids):
    result=client.post(BASE+'/counts',json=dict(factory_id='huakang-c',request_id=str(uuid4()),batch_ids=ids,reason='现场盘点'))
    assert result.status_code==200,result.text
    return result.json()


def act(client, doc, action='POST', actual='7.000001', **changes):
    body=dict(factory_id='huakang-c',request_id=str(uuid4()),expected_revision=doc['revision'],action=action,business_date=date.today().isoformat(),confirmed=True,reason='核对实物',
        lines=[dict(batch_id=l['batch_id'],actual=actual,reason='核实差异') for l in doc['lines']])
    body.update(changes)
    return client.post(BASE+'/counts/'+doc['id'],json=body),body


def test_opening_template_atomic_post_repeat_precision_and_reports(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client,'admin')
        template=client.get(BASE+'/template',params={'factory_id':'huakang-c'})
        assert template.status_code==200
        workbook=load_workbook(io.BytesIO(template.content));assert workbook.active['B2'].number_format=='@'
        assert '可用仓位' in workbook.sheetnames
        content=book([line(),line(**{'来源行号':'002','物料编码':'00013','物料名称':'扣子','物料分类':'辅料','单位':'个','缸号':'','卷号':'','期初数量':'3'})])
        result=preview(client,content).json();assert result['new_count']==2 and result['totals']=={'码':'10.000001','个':'3'}
        assert workspace(client)['stock']==[]
        request=str(uuid4());posted=apply(client,content,result,request);assert posted.status_code==200,posted.text
        assert apply(client,content,result,request).json()==posted.json()
        assert apply(client,content,result,request,source_no='另一份').status_code==409
        assert len(workspace(client)['stock'])==2
        repeat=preview(client,content).json();assert repeat['skip_count']==2
        again=apply(client,content,repeat);assert again.status_code==200 and again.json()['posted_count']==0
        altered=preview(client,book([line(**{'期初数量':'11'})])).json();assert altered['error_count']==1
        report=client.get('/api/warehouse-operations/fabric/order-summary',params={'factory_id':'huakang-c'}).json()
        assert {r['unit']:r['balance'] for r in report['rows']}=={'码':'10.000001','个':'3'}
        monthly=client.get('/api/warehouse-operations/fabric/monthly',params={'factory_id':'huakang-c','month':'2026-07'}).json()
        assert monthly['details']==[]
        assert not any(r['kind']=='RECEIPT' for r in workspace(client)['documents'])


def test_opening_invalid_rows_and_stale_preview_never_partially_post(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client,'admin')
        for changed in [{'期初数量':'0'},{'期初数量':'1.0000001'},{'物料编码':12},{'期初数量':'=SUM(1,2)'},{'仓位':'不存在'},{'缸号':''}]:
            content=book([line(),line(**{'来源行号':'002','物料编码':'00013',**changed})])
            result=preview(client,content).json();assert result['error_count']>0
            assert apply(client,content,result).status_code==422
            assert workspace(client)['stock']==[]
        content=book();result=preview(client,content).json()
        existing=post(client,payload(client,'RECEIPT',lot='别的缸号'))
        assert apply(client,content,result).status_code==409
        assert len(workspace(client)['stock'])==1
        collision=preview(client,book([line(**{'缸号':'别的缸号','卷号':''})])).json()
        assert collision['error_count']==1
        assert existing['id']==workspace(client)['stock'][0]['id']


def test_count_drafts_zero_variance_replay_and_prior_issue_correction_protection(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client,'admin')
        receipt=post(client,payload(client,'RECEIPT'))
        issue=post(client,payload(client,'ISSUE',batch_id=receipt['id'],quantity='3'))
        doc=count(client,[receipt['id']])
        saved,_=act(client,doc,'SAVE',actual='');assert saved.status_code==200,saved.text
        assert workspace(client)['stock'][0]['quantity']=='7'
        doc=saved.json()
        empty,_=act(client,doc,actual='');assert empty.status_code==422
        mismatch,_=act(client,doc,lines=[]);assert mismatch.status_code==422
        posted,body=act(client,doc,actual='7');assert posted.status_code==200,posted.text
        assert client.post(BASE+'/counts/'+doc['id'],json=body).json()==posted.json()
        assert act(client,doc,actual='7')[0].status_code==409
        post(client,payload(client,'CORRECTION',original_id=issue['id'],business_date=date.today().isoformat()),status=409)
        assert workspace(client)['stock'][0]['quantity']=='7'
        # Generic operation endpoint cannot bypass the specialized posting authority.
        post(client,payload(client,'STOCKTAKE',batch_id=receipt['id'],quantity='20'),status=422)


def test_count_stale_same_quantity_refresh_gain_loss_and_restrictions(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client,'admin')
        receipt=post(client,payload(client,'RECEIPT'))
        doc=count(client,[receipt['id']])
        issue=post(client,payload(client,'ISSUE',batch_id=receipt['id'],quantity='2'))
        post(client,payload(client,'CORRECTION',original_id=issue['id']))
        assert workspace(client)['stock'][0]['quantity']=='10'
        assert act(client,doc,actual='10')[0].status_code==409
        refreshed,_=act(client,doc,'REFRESH');assert refreshed.status_code==200
        doc=refreshed.json();assert doc['lines'][0]['actual']==''
        missing,_=act(client,doc,actual='12',lines=[dict(batch_id=receipt['id'],actual='12',reason='')]);assert missing.status_code==422
        posted,_=act(client,doc,actual='12.000001');assert posted.status_code==200,posted.text
        assert workspace(client)['stock'][0]['quantity']=='12.000001'
        next_doc=count(client,[receipt['id']])
        zero,_=act(client,next_doc,actual='0');assert zero.status_code==200,zero.text
        assert Decimal(workspace(client)['stock'][0]['quantity'])==0
        # An independent restricted batch retains its QC restriction after a gain.
        held=post(client,payload(client,'RECEIPT',item_code='00013'))
        quality(client,held['id'],state='HOLD')
        held_doc=count(client,[held['id']]);assert act(client,held_doc,actual='15')[0].status_code==200
        held_batch=next(b for b in workspace(client)['stock'] if b['id']==held['id'])
        assert held_batch['quantity']=='15' and held_batch['available_quantity']=='0' and held_batch['quality_status']=='HOLD'
        report=client.get('/api/warehouse-operations/fabric/order-summary',params={'factory_id':'huakang-c'}).json()
        assert sum(Decimal(row['balance']) for row in report['rows'])==15


def test_count_transfer_and_original_receipt_correction_dependencies(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client,'admin')
        receipt=post(client,payload(client,'RECEIPT'))
        move=post(client,payload(client,'TRANSFER',batch_id=receipt['id'],quantity='4',location='A02'))
        doc=count(client,[receipt['id']]);assert act(client,doc,actual='6')[0].status_code==200
        post(client,payload(client,'CORRECTION',original_id=move['id'],business_date=date.today().isoformat()),status=409)
        from test_fabric_receipt_corrections import original,reverse
        _,original_receipt,_,_=original(client)
        doc=count(client,[original_receipt['batches'][0]['id']]);assert act(client,doc,actual=doc['lines'][0]['snapshot']['quantity'])[0].status_code==200
        response,_=reverse(client,original_receipt,business_date=date.today().isoformat())
        assert response.status_code==409,response.text


def test_count_cancel_dates_and_factory_access(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client,'admin')
        receipt=post(client,payload(client,'RECEIPT'))
        doc=count(client,[receipt['id']])
        response,_=act(client,doc,business_date='2026-06-01');assert response.status_code==422
        cancelled,_=act(client,doc,'CANCEL');assert cancelled.status_code==200
        assert act(client,cancelled.json(),'POST')[0].status_code==409
        assert workspace(client)['stock'][0]['quantity']=='10'
        for path in [BASE,BASE+'/template']:
            assert client.get(path,params={'factory_id':'huakang-b'}).status_code in (403,422)
        assert client.post(BASE+'/counts',json=dict(factory_id='huakang-b',request_id=str(uuid4()),batch_ids=[receipt['id']],reason='x')).status_code==422


def test_opening_rejects_cells_outside_template_bounds():
    from app.services.fabric_inventory_setup import parse
    from fastapi import HTTPException
    for address in ('B1003', 'N2'):
        workbook = load_workbook(io.BytesIO(book()))
        workbook.active[address] = '不能丢掉的库存资料'
        output = io.BytesIO(); workbook.save(output)
        with pytest.raises(HTTPException) as caught:
            parse('opening.xlsx', output.getvalue())
        assert caught.value.status_code == 422
        assert '标准列范围' in caught.value.detail


def test_template_exports_location_identity_as_text_not_formulas(monkeypatch):
    from app.services import fabric_inventory_setup as setup
    monkeypatch.setattr(setup.ops.locations, 'catalog', lambda *_: [dict(status='ACTIVE', warehouse='=1+1', code='00001')])
    workbook = load_workbook(io.BytesIO(setup.template(None)))
    assert workbook['可用仓位']['A2'].value == '=1+1'
    assert workbook['可用仓位']['A2'].data_type == 's'
    assert workbook['可用仓位']['B2'].value == '00001'
    assert workbook['可用仓位']['B2'].number_format == '@'


def test_posting_permissions_and_atomic_failure_of_later_count_line(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, 'admin')
        first = post(client, payload(client, 'RECEIPT'))
        second = post(client, payload(client, 'RECEIPT', item_code='00013'))
        doc = count(client, [first['id'], second['id']])
        saved, _ = act(client, doc, 'SAVE', actual='8')
        doc = saved.json()
        api = importlib.import_module('app.api.warehouse_operations')
        with monkeypatch.context() as scope:
            scope.setattr(api, 'can', lambda actor, permission, *args: permission in {'fabric_operations:read', 'fabric_operations:operate'})
            assert act(client, doc)[0].status_code == 403
            assert preview(client, book()).status_code == 403
        with monkeypatch.context() as scope:
            scope.setattr(api, 'can', lambda actor, permission, *args: permission in {'fabric_operations:read', 'fabric_operations:correct'})
            assert act(client, doc, 'SAVE')[0].status_code == 403
            lines = [dict(batch_id=first['id'], actual='8', reason='盘亏'), dict(batch_id=second['id'], actual='9', reason='')]
            assert act(client, doc, lines=lines)[0].status_code == 422
            assert all(row['quantity'] == '10' for row in workspace(client)['stock'])
            assert not any(row['kind'] == 'STOCKTAKE' for row in workspace(client)['documents'])
            assert client.get(BASE, params={'factory_id': 'huakang-c'}).json()['counts'][0]['revision'] == doc['revision']
            posted, _ = act(client, doc, actual='8')
            assert posted.status_code == 200, posted.text
            assert all(row['quantity'] == '8' for row in workspace(client)['stock'])


def test_concurrent_count_confirmations_only_post_once(monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    with make_client(monkeypatch) as client:
        login_as(client, 'admin')
        receipt = post(client, payload(client, 'RECEIPT'))
        doc = count(client, [receipt['id']])
        barrier = Barrier(2)
        def submit(_):
            barrier.wait(timeout=10)
            return act(client, doc, actual='8')[0].status_code
        with ThreadPoolExecutor(max_workers=2) as pool:
            assert sorted(pool.map(submit, range(2))) == [200, 409]
        assert workspace(client)['stock'][0]['quantity'] == '8'


def test_count_ignores_other_batches_but_exposes_exact_changed_batch_ids(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, 'admin')
        first = post(client, payload(client, 'RECEIPT'))
        second = post(client, payload(client, 'RECEIPT', lot='00002'))
        sibling = post(client, payload(client, 'TRANSFER', batch_id=first['id'], quantity='2', location='A02'))
        doc = count(client, [first['id']])
        other = post(client, payload(client, 'ISSUE', batch_id=second['id'], quantity='2'))
        post(client, payload(client, 'CORRECTION', original_id=other['id']))
        # Neither the same material nor a moved sibling expands the selected scope.
        post(client, payload(client, 'ISSUE', batch_id=sibling['id'], quantity='1'))
        post(client, payload(client, 'RECEIPT', item_code='00014'))
        current = client.get(BASE, params={'factory_id': 'huakang-c'}).json()['counts'][0]
        assert not current['stale'] and current['changed_batch_ids'] == []
        assert act(client, doc, actual='9')[0].status_code == 200
        # A different zero-variance count of the same batch is still a relevant event.
        left = count(client, [first['id'], second['id']])
        right = count(client, [second['id']])
        assert act(client, right, actual='10')[0].status_code == 200
        current = next(r for r in client.get(BASE, params={'factory_id': 'huakang-c'}).json()['counts'] if r['id'] == left['id'])
        assert current['stale'] and current['changed_batch_ids'] == [second['id']]
        assert act(client, left)[0].status_code == 409


@pytest.mark.parametrize('change', ['return', 'transfer_reverse', 'transfer_target_reverse', 'quality', 'original_reverse'])
def test_selected_batch_related_evidence_invalidates_count(monkeypatch, change):
    with make_client(monkeypatch) as client:
        login_as(client, 'admin')
        if change == 'original_reverse':
            from test_fabric_receipt_corrections import original, reverse
            _, receipt, _, _ = original(client)
            doc = count(client, [receipt['batches'][0]['id']])
            assert reverse(client, receipt)[0].status_code == 200
        else:
            receipt = post(client, payload(client, 'RECEIPT'))
            if change == 'return':
                issue = post(client, payload(client, 'ISSUE', batch_id=receipt['id'], quantity='2'))
                doc = count(client, [receipt['id']])
                post(client, payload(client, 'RETURN', original_id=issue['id'], quantity='1'))
            elif change.startswith('transfer'):
                move = post(client, payload(client, 'TRANSFER', batch_id=receipt['id'], location='A02', quantity='2'))
                doc = count(client, [move['id'] if change == 'transfer_target_reverse' else receipt['id']])
                post(client, payload(client, 'CORRECTION', original_id=move['id']))
            else:
                doc = count(client, [receipt['id']])
                quality(client, receipt['id'], state='HOLD')
        current = client.get(BASE, params={'factory_id': 'huakang-c'}).json()['counts'][0]
        assert current['stale'] and current['changed_batch_ids'] == [doc['lines'][0]['batch_id']]
        assert act(client, doc)[0].status_code == 409


def test_old_count_draft_remains_conservative_until_explicit_refresh(monkeypatch):
    import json
    with make_client(monkeypatch) as client:
        login_as(client, 'admin')
        first = post(client, payload(client, 'RECEIPT'))
        doc = count(client, [first['id']])
        database = importlib.import_module('app.db')
        model = importlib.import_module('app.models.warehouse_operations').WarehouseOperation
        with database.SessionLocal() as db:
            row = db.get(model, doc['id'])
            body = json.loads(row.payload_json)
            body['result'].pop('count_scope'); body['result'].pop('batch_tokens')
            row.payload_json = json.dumps(body); db.commit()
        post(client, payload(client, 'RECEIPT', item_code='00013'))
        current = client.get(BASE, params={'factory_id': 'huakang-c'}).json()['counts'][0]
        assert current['stale'] and current['legacy_scope']
        assert act(client, doc)[0].status_code == 409
        refreshed, _ = act(client, doc, 'REFRESH')
        assert refreshed.status_code == 200
        doc = refreshed.json()
        assert doc['count_scope'] == 'BATCHES_V1' and doc['lines'][0]['actual'] == ''
        post(client, payload(client, 'RECEIPT', item_code='00014'))
        assert act(client, doc)[0].status_code == 200
        assert len([row for row in workspace(client)['documents'] if row['kind'] == 'STOCKTAKE']) == 1
