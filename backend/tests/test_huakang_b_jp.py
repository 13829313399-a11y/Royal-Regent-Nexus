"""Disposable fixtures only: no desktop app startup, real mail or price writes."""
from copy import deepcopy
from datetime import date
from email.message import EmailMessage
from io import BytesIO
import importlib.util
import json
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile
import asyncio
import openpyxl
import pytest
from alembic.migration import MigrationContext
from alembic.operations import Operations
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event, select, func, text
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from app.api import huakang_b_jp as api
from app.db import get_db
from app.services.auth import get_current_user
from app.services.huakang_b_jp import intake, parser
from app.models.huakang_b_jp import JPReference, JPReferenceEvent, JPCounter, JPExport
from app.models.customer_order_ledger import OrderLedgerLine as Line, OrderLedgerSource as Source
from test_customer_order_ledger import user, migration as ledger_migration


def migration():
    path = Path(__file__).parents[1] / "alembic/legacy/20261015_0163_huakang_b_jp_intake.py"
    spec = importlib.util.spec_from_file_location("jp_migration", path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


@pytest.fixture
def engine():
    engine = create_engine("sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False})
    @event.listens_for(engine, "connect")
    def fk(conn, _): conn.execute("PRAGMA foreign_keys=ON")
    with engine.begin() as conn:
        with Operations.context(MigrationContext.configure(conn)):
            ledger_migration().upgrade()
            migration().upgrade()
    yield engine
    engine.dispose()


@pytest.fixture
def db(engine):
    with Session(engine, expire_on_commit=False, autoflush=False) as db:
        yield db


@pytest.fixture
def client(db, monkeypatch):
    from app.services import business_authz
    monkeypatch.setattr(business_authz.settings, "authz_mode", "enforce")
    monkeypatch.setattr(business_authz, "add_auth_audit", lambda *a, **k: None)
    app = FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: user(["read", "write", "export"], factory="huakang-b")
    monkeypatch.setattr(api, "refresh_user", lambda req, session: app.dependency_overrides[get_current_user]())
    with TestClient(app) as client:
        client.headers["X-JP-Intake"] = "1"
        client.test_app = app
        yield client


def xlsx(kind):
    w = openpyxl.Workbook()
    s = w.active
    if kind == "schedule":
        s.title = "华康B"
        s.append([]); s.append([]); s.append(["订单编号"] * 15)
        # Same item but different customer/PO so fixture mail is not a duplicate.
        s.append(["0126-21614", "OTHER", "0000000001", "", "", "67053", "67053 WB0", "67053", "Test toy", "3L", 100, "", 3, "美国", ""])
    elif kind == "price":
        s.append(["货号编码","产品货号","客货价(HK$)","货运类别","长货号","货价日期","外箱-装箱数"])
        s.append(["CODE","67053",12.5,"柜货","67053-000-2B-003-WB0",date(2026,1,1),3])
        s.append(["CODE2","67053",10,"出厂价","67053-000-2B-003-WB0",date(2026,1,1),3])
    else:
        s.append(["货号","邮件长货号","货运类别","装箱数","货币","候选客货价(HK$)","确认价格","确认币种"])
        s.append(["67053","67053-000-2B-003-WB0","柜货",3,"USD",23.4,3,"USD"])
    out = BytesIO(); w.save(out)
    return out.getvalue()


def body(qty="648", header=True, po="0117978388"):
    headings = "<tr><th>Item</th><th>Qty</th><th>Customer PO</th></tr>" if header else ""
    return f"<p>S/D 12/31/2026</p><table>{headings}<tr><td>67053-000-2B-003-WB0</td><td>{qty}</td><td>{po}</td></tr></table>"


def mail(html=None, subject="WMUS order inquiry", stamp="Tue, 13 Oct 2026 10:00:00 +0800"):
    m = EmailMessage()
    m["Subject"] = subject; m["Date"] = stamp
    m.set_content("test")
    m.add_alternative(html or body(), subtype="html")
    return m.as_bytes()


def install(db):
    for kind in ("schedule", "price"):
        content = xlsx(kind)
        intake._jp_validate_reference(kind, content)
        api.transaction(db, lambda: intake.save_reference(db, kind, kind + ".xlsx", content, 0, "业务员"))


def preview(db, files=None):
    return intake._jp_preview(files or [("test.eml", mail())], "", "", "2026-10-13", intake.snapshots(db), [], "u1")


def export(db, p, overrides=None, files=None):
    return api.transaction(db, lambda: intake.save_export(db, p, [r["id"] for r in p["rows"]], overrides or [],
        intake.snapshots(db), files or [("test.eml", mail())], "", "逐项核对正确", "业务员"))


def test_single_row_headerless_is_supported_and_numeric_ambiguity_not_guessed():
    parsed = parser.parse_order_email(body(header=False), "WMUS", date(2026,10,13))
    assert len(parsed["orders"]) == 1 and parsed["orders"][0]["qty"] == "648"
    html = "<table><tr><td>67053-000-2B-003-WB0</td><td>25</td><td>648</td></tr></table>"
    with pytest.raises(parser.QuantityColumnRequired):
        parser.parse_order_email(html, "WMUS", date(2026,10,13))
    corrected = parser.parse_order_email(html, "WMUS", date(2026,10,13), quantity_columns={0: 2})
    assert corrected["orders"][0]["qty"] == "648"


@pytest.mark.parametrize('header', [False, True])
def test_explicit_current_reply_can_update_the_nearest_quoted_order(header):
    html = '<p>DDD 9/18, 货期1/22. 8L pkg.</p><p>发件人：旧回复</p>' + body(header=header)
    parsed = parser.parse_order_email(html, 'WMUS order inquiry', date(2026, 9, 10))
    assert len(parsed['orders']) == 1
    row = parsed['orders'][0]
    assert row['ship_date'] == '2027/1/15' and row['cutoff_date'] == '2027/1/22'
    assert row['pkg'] == '8L'
    canceled = parser.parse_order_email('<p>Cancel this order. 货期1/22.</p><p>发件人：旧回复</p>' + body(header=header),
                                        'WMUS order inquiry', date(2026, 9, 10))
    assert not canceled['orders']


def test_quoted_table_values_cannot_override_the_explicit_current_reply_update():
    table = ('<table><tr><th>Item</th><th>Qty</th><th>Customer PO</th><th>Promised CRD</th><th>Pkg</th></tr>'
             '<tr><td>67053-000-2B-003-WB0</td><td>648</td><td>0117978388</td><td>12/31/2026</td><td>3L</td></tr></table>')
    reply = '<p>货期1/22. 8L pkg.</p>'
    parsed = parser.parse_order_email(reply + '<p>发件人：旧回复</p>' + table, 'WMUS', date(2026, 9, 10))
    row = parsed['orders'][0]
    assert (row['cutoff_date'], row['ship_date'], row['pkg']) == ('2027/1/22', '2027/1/15', '8L')
    assert row['quoted_original'] == {'cutoff_date': '2026/12/31', 'ship_date': '2026/12/24', 'pkg': '3L'}
    assert any('引用表货期' in w for w in parsed['warnings'])
    assert any('引用表包装' in w for w in parsed['warnings'])
    # The established current-table precedence is deliberately unchanged.
    current = parser.parse_order_email(reply + table, 'WMUS', date(2026, 9, 10))['orders'][0]
    assert current['cutoff_date'] == '2026/12/31' and current['pkg'] == '3L'
    assert 'quoted_original' not in current


@pytest.mark.parametrize('header', [False, True])
def test_quoted_reply_update_survives_schedule_enrichment_and_export(db, header):
    install(db)
    html = ('<p>货期1/22. 8L pkg.</p><p>发件人：旧回复</p><table>'
            '<tr><th>Item</th><th>Qty</th><th>Customer PO</th><th>Promised CRD</th><th>Pkg</th></tr>'
            '<tr><td>67053-000-2B-003-WB0</td><td>648</td><td>0117978388</td><td>12/31/2026</td><td>3L</td></tr></table>')
    if not header:
        html = '<p>货期1/22. 8L pkg.</p><p>发件人：旧回复</p>' + body(header=False)
    files = [('reply.eml', mail(html))]
    p = preview(db, files)
    row = p['rows'][0]
    assert row['pkg'] == '8L' and row['cutoff_date'] == '2027/1/22'
    assert any('最新回复包装 8L 与排期历史 3L 不同' in w for w in p['warnings'])
    record = export(db, p, files=files)
    assert record.data['rows'][0]['pkg'] == '8L'
    if header:
        assert record.data['rows'][0]['quoted_original']['pkg'] == '3L'
    assert record.data['rows'][0]['quoted_original']


def test_preview_enrichment_and_44_column_export_preserve_originals(db):
    install(db)
    refs = deepcopy(intake.reference_state(db))
    files = [("test.eml", mail())]
    original = deepcopy(files)
    p = preview(db, files)
    row = p["rows"][0]
    assert row["qty"] == "648"
    assert row["pkg"] == "3L" and row["price_t"] == "12.5" and row["factory_price_t"] == "10"
    assert row["ship_date"] == "2026/12/24" and row["cutoff_date"] == "2026/12/31"
    assert not row["exists"]
    record = export(db, p)
    w = openpyxl.load_workbook(BytesIO(record.content))
    s = w.active
    assert s.max_column == 44 and s.max_row == 2
    assert s.cell(2,1).value == "0126-21615"
    assert s.cell(2,3).value == "0117978388" and s.cell(2,3).data_type == "s"
    assert s.cell(2,13).value == 648 and s.cell(2,21).value == 12.5
    assert s.cell(2,37).value == "问单" and s.cell(2,38).value == "华康B"
    from datetime import datetime
    assert s.cell(2,18).value == datetime(2026,12,24)
    assert s.cell(2,18).number_format == "yyyy/mm/dd"
    assert intake.reference_state(db) == refs and files == original
    assert db.scalar(select(func.count()).select_from(Line)) == 0


def test_invalid_reference_or_stale_upload_preserves_old_version(client, db):
    install(db)
    before = intake.reference_state(db)
    result = client.post("/api/huakang-b-jp/references/schedule", data={"factory_id":"huakang-b","expected_revision":1},
        files={"file":("invalid.xlsx",b"bad","application/octet-stream")})
    assert result.status_code == 400
    assert intake.reference_state(db) == before
    result = client.post("/api/huakang-b-jp/references/schedule", data={"factory_id":"huakang-b","expected_revision":0},
        files={"file":("same.xlsx",xlsx("schedule"),"application/octet-stream")})
    assert result.status_code == 409
    assert intake.reference_state(db) == before


def test_numbering_idempotency_and_later_batch(db):
    install(db)
    p = preview(db)
    first = export(db, p)
    again = export(db, p)
    assert first.id == again.id
    assert db.get(JPCounter, "huakang-b").last_number == 21615
    files = [("next.eml", mail(body(po="0117978389"), subject="Other customer"))]
    second = export(db, preview(db, files), files=files)
    assert second.first_number == 21616
    assert db.scalar(select(func.count()).select_from(JPExport)) == 2
    # Even a lost counter cannot reuse reserved numbers retained in evidence.
    db.execute(text("DELETE FROM jp_intake_counters")); db.commit()
    files = [("third.eml", mail(body(po="0117978390"), subject="Third customer"))]
    third = export(db, preview(db, files), files=files)
    assert third.first_number == 21617


def test_fresh_reference_between_preview_and_commit_rejects(db):
    install(db)
    p = preview(db)
    api.transaction(db, lambda: intake.save_reference(db, "confirmed", "confirmed.xlsx", xlsx("confirmed"), 0, "文员"))
    with pytest.raises(HTTPException) as exc:
        export(db, p)
    assert exc.value.status_code == 409
    assert db.get(JPCounter, "huakang-b") is None
    assert db.scalar(select(func.count()).select_from(JPExport)) == 0


@pytest.mark.parametrize("qty", ["-5", "0", "1.5", "NaN", "100000000000"])
def test_bad_quantities_never_allocate_numbers(db, qty):
    install(db)
    files = [("bad.eml", mail(body(qty=qty)))]
    p = preview(db, files)
    with pytest.raises(HTTPException):
        export(db, p, files=files)
    assert db.get(JPCounter, "huakang-b") is None


def test_formula_strings_are_text_not_excel_formulas(db):
    install(db)
    p = preview(db)
    change = [{"id":p["rows"][0]["id"],"field":"note","value":"=HYPERLINK(\"https://example.invalid\")"}]
    rec = export(db, p, overrides=change)
    w = openpyxl.load_workbook(BytesIO(rec.content))
    assert w.active.cell(2,28).data_type == "s"


def test_permissions_and_factory_isolation(client, db):
    install(db)
    assert client.get("/api/huakang-b-jp/references?factory_id=huakang-a").status_code == 400
    client.test_app.dependency_overrides[get_current_user] = lambda: user(["read"], factory="huakang-a")
    assert client.get("/api/huakang-b-jp/references?factory_id=huakang-b").status_code == 403
    client.test_app.dependency_overrides[get_current_user] = lambda: user(["read"], factory="huakang-b")
    assert client.get("/api/huakang-b-jp/references?factory_id=huakang-b").status_code == 200
    assert client.post("/api/huakang-b-jp/references/price", data={"factory_id":"huakang-b","expected_revision":1},
        files={"file":("prices.xlsx",xlsx("price"),"application/octet-stream")}).status_code == 403
    assert client.post("/api/huakang-b-jp/export", data={"factory_id":"huakang-b","received_date":"2026-10-13",
        "fingerprint":"x","selected":"[]","reason":"核对正确","confirmed":True}).status_code == 403


def test_real_worker_http_preview_export_download_and_stale_fingerprint(client, db):
    install(db)
    form = {"factory_id":"huakang-b","received_date":"2026-10-13","merchandiser":"测试跟单"}
    files = [("files",("test.eml",mail(),"message/rfc822"))]
    response = client.post("/api/huakang-b-jp/preview", data=form, files=files)
    assert response.status_code == 200, response.text
    p = response.json()
    form.update(fingerprint=p["fingerprint"], selected=json.dumps([r["id"] for r in p["rows"]]),
                reason="逐项核对正确", confirmed=True)
    response = client.post("/api/huakang-b-jp/export", data=form, files=files)
    assert response.status_code == 200, response.text
    id = response.json()["id"]
    again = client.post("/api/huakang-b-jp/export", data=form, files=files)
    assert again.json()["id"] == id
    content = client.get(f"/api/huakang-b-jp/exports/{id}/download?factory_id=huakang-b")
    assert content.status_code == 200 and len(content.content) > 100
    assert content.headers["cache-control"] == "no-store"
    form["fingerprint"] = "wrong"
    assert client.post("/api/huakang-b-jp/export", data=form, files=files).status_code == 409


def test_export_evidence_is_immutable_and_downgrade_refused(db):
    install(db)
    rec = export(db, preview(db))
    with pytest.raises(Exception):
        db.execute(text("DELETE FROM jp_intake_exports WHERE id=:id"), {"id":rec.id}); db.commit()
    db.rollback()
    with pytest.raises(RuntimeError):
        with Operations.context(MigrationContext.configure(db.connection())):
            migration().downgrade()
    db.rollback()


def test_migration_is_additive_and_empty_downgrade_preserves_other_tables(db):
    from sqlalchemy import inspect
    db.execute(text('CREATE TABLE unrelated_module_sentinel (value INTEGER)'))
    db.execute(text('INSERT INTO unrelated_module_sentinel VALUES (123)')); db.commit()
    before = set(inspect(db.connection()).get_table_names())
    with Operations.context(MigrationContext.configure(db.connection())):
        migration().downgrade()
    assert set(inspect(db.connection()).get_table_names()) == before - set(migration().NAMES)
    assert db.scalar(text('SELECT value FROM unrelated_module_sentinel')) == 123
    with Operations.context(MigrationContext.configure(db.connection())):
        migration().upgrade()
    assert set(inspect(db.connection()).get_table_names()) == before


def test_postgres_offline_migration_creates_one_guard_and_five_tables():
    from io import StringIO
    out = StringIO()
    context = MigrationContext.configure(dialect_name='postgresql', opts={'as_sql': True, 'output_buffer': out})
    with Operations.context(context): migration().upgrade()
    sql = out.getvalue()
    assert sql.count('CREATE TABLE jp_intake_') == 5
    assert sql.count('CREATE FUNCTION jp_intake_immutable') == 1
    assert sql.count('CREATE TRIGGER ') == 3
    assert 'ALTER TABLE order_ledger' not in sql and 'DROP TABLE' not in sql


def test_zip_caps_and_empty_or_quoted_only_orders_fail_closed(db):
    install(db)
    with BytesIO() as buffer:
        with ZipFile(buffer,"w") as z:
            for i in range(201): z.writestr(f"{i}.eml", b"x")
        with pytest.raises(ValueError):
            intake.expand([("bad.zip", buffer.getvalue())])
    files = [("old.eml", mail("<p>No new orders</p><p>Original Message</p>" + body(header=False)))]
    p = preview(db, files)
    assert not p["rows"] and p["failed_sources"]
    with pytest.raises(HTTPException):
        export(db, p, files=files)


def test_cross_origin_form_header_missing_is_rejected(client, db):
    install(db)
    del client.headers["X-JP-Intake"]
    response = client.post("/api/huakang-b-jp/preview", data={"factory_id":"huakang-b","received_date":"2026-10-13"},
        files=[("files",("test.eml",mail(),"message/rfc822"))])
    assert response.status_code == 403


def test_permissions_revoked_during_worker_prevent_write(client, db, monkeypatch):
    install(db)
    original = api.run_order_job
    async def revoke(*args, **kwargs):
        result = await original(*args, **kwargs)
        client.test_app.dependency_overrides[get_current_user] = lambda: user(["read"], factory="huakang-b")
        return result
    monkeypatch.setattr(api, "run_order_job", revoke)
    before = intake.reference_state(db)
    response = client.post("/api/huakang-b-jp/references/confirmed", data={"factory_id":"huakang-b","expected_revision":0},
        files={"file":("confirmed.xlsx",xlsx("confirmed"),"application/octet-stream")})
    assert response.status_code == 403
    assert intake.reference_state(db) == before


def test_parallel_exports_do_not_reuse_numbers(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    engine = create_engine("sqlite:///" + str(tmp_path / "jp.db"), connect_args={"timeout":10,"check_same_thread":False})
    with engine.begin() as conn:
        with Operations.context(MigrationContext.configure(conn)):
            ledger_migration().upgrade(); migration().upgrade()
    with Session(engine, expire_on_commit=False) as db:
        install(db)
        refs = intake.snapshots(db)
    files = [[(f"{i}.eml", mail(body(po=f"011797838{i}"), subject=f"Buyer {i}"))] for i in range(2)]
    previews = [intake._jp_preview(f, "", "", "2026-10-13", refs, [], "u1") for f in files]
    def run(i):
        with Session(engine, expire_on_commit=False) as db:
            rec = api.transaction(db, lambda: intake.save_export(db, previews[i], [r["id"] for r in previews[i]["rows"]],
                [], refs, files[i], "", "逐项核对正确", "业务员"))
            return rec.first_number
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(run, range(2)))
    assert sorted(results) == [21615,21616]
    engine.dispose()


def test_same_business_line_by_another_filename_or_actor_cannot_reserve_again(db):
    install(db)
    first = export(db, preview(db))
    files = [("renamed.eml", mail())]
    p = intake._jp_preview(files, "", "", "2026-10-13", intake.snapshots(db), [], "another-user")
    with pytest.raises(HTTPException) as exc:
        export(db, p, files=files)
    assert exc.value.status_code == 409
    assert db.get(JPCounter, "huakang-b").last_number == first.last_number


def test_same_subject_different_pos_are_not_treated_as_revisions(db):
    install(db)
    files = [('one.eml', mail(body(po='1111111111'), stamp='Tue, 13 Oct 2026 10:00:00 +0800')),
             ('two.eml', mail(body(po='2222222222'), stamp='Tue, 13 Oct 2026 11:00:00 +0800'))]
    p = preview(db, files)
    assert {r['customer_po'] for r in p['rows']} == {'1111111111', '2222222222'}
    assert not p['failed_sources']
    assert export(db, p, files=files).last_number == 21616


def test_exact_po_revision_keeps_newest_but_unknown_po_needs_review(db):
    install(db)
    files = [('old.eml', mail(body(qty='100'), stamp='Tue, 13 Oct 2026 10:00:00 +0800')),
             ('new.eml', mail(body(qty='200'), stamp='Tue, 13 Oct 2026 11:00:00 +0800'))]
    p = preview(db, files)
    assert len(p['rows']) == 1 and p['rows'][0]['qty'] == '200'
    files = [('old.eml', mail(body(qty='100', po='TBC'), stamp='Tue, 13 Oct 2026 10:00:00 +0800')),
             ('new.eml', mail(body(qty='200', po='TBC'), stamp='Tue, 13 Oct 2026 11:00:00 +0800'))]
    p = preview(db, files)
    assert p['failed_sources']
    with pytest.raises(HTTPException): export(db, p, files=files)
    same = mail()
    p = preview(db, [('same.eml', same), ('renamed.eml', same)])
    assert len(p['rows']) == 1 and not p['failed_sources']


@pytest.mark.parametrize('field,value', [('customer', 'Amazon US'), ('carton_t', '6')])
def test_changed_price_context_requires_explicit_price_review(db, field, value):
    install(db)
    p = preview(db)
    row = p['rows'][0]
    changed = [dict(id=row['id'], field=field, value=value)]
    with pytest.raises(HTTPException, match='重新填写核对货价'):
        export(db, p, overrides=changed)
    assert db.get(JPCounter, 'huakang-b') is None
    changed += [dict(id=row['id'], field='price_t', value='20'),
                dict(id=row['id'], field='factory_price_t', value='')]
    record = export(db, p, overrides=changed)
    assert record.data['rows'][0]['manual_price_review']
    assert openpyxl.load_workbook(BytesIO(record.content)).active.cell(2, 21).value == 20


def test_low_number_output_keeps_four_digit_minimum_format():
    assert intake.workbook([dict(customer="Buyer", customer_po="0001", order_date="2026/10/13",
        baby_item="67053", qty="3", carton_t="3", price_t="1", factory_price_t="", ship_date="2026/12/24", cutoff_date="2026/12/31")],
        1, "0126", "业务员")
    data = intake.workbook([dict(customer="Buyer", customer_po="0001", order_date="2026/10/13",
        baby_item="67053", qty="3", carton_t="3", price_t="1", factory_price_t="", ship_date="2026/12/24", cutoff_date="2026/12/31")],
        1, "0126", "业务员")
    assert openpyxl.load_workbook(BytesIO(data)).active.cell(2,1).value == "0126-0001"


def test_confirmed_prices_can_be_disabled_without_losing_history(client, db):
    install(db)
    api.transaction(db, lambda: intake.save_reference(db, "confirmed", "confirmed.xlsx", xlsx("confirmed"), 0, "文员"))
    before = db.scalar(select(func.count()).select_from(JPReferenceEvent))
    response = client.post("/api/huakang-b-jp/references/confirmed/clear", data={"factory_id":"huakang-b","expected_revision":1})
    assert response.status_code == 200
    assert response.json()["confirmed"]["revision"] == 2
    assert db.scalar(select(func.count()).select_from(JPReferenceEvent)) == before + 1
    assert preview(db)["rows"]


def confirmation(amount, currency=""):
    w = openpyxl.Workbook()
    w.active.append(["产品货号", "确认价格", "确认币种", "货币", "候选客货价(HK$)"])
    w.active.append(["67053", amount, currency, "USD", 23.4])
    out = BytesIO(); w.save(out)
    return out.getvalue()


def test_confirmation_requires_explicit_currency_and_versions_upload_choice(db, tmp_path):
    from app.services.huakang_b_jp.prices import PriceTable
    content = confirmation(23.4)
    with pytest.raises(ValueError, match="未明确币种"):
        intake._jp_validate_reference("confirmed", content)
    assert intake._jp_validate_reference("confirmed", content, {"currency":"USD"}) == {"rows":1}
    path = tmp_path / "confirmed.xlsx"; path.write_bytes(content)
    obj = PriceTable(); obj.import_confirmed_xlsx(path, default_currency="USD")
    assert next(iter(obj.confirmed.values()))["price"] == "182.52"  # Never guess HKD from the candidate amount.
    for revision, currency in ((0, "USD"), (1, "HKD")):
        result = api.transaction(db, lambda: intake.save_reference(db, "confirmed", "confirmed.xlsx", content,
            revision, "文员", {"currency":currency}))
        assert result["confirmed"]["revision"] == revision + 1
    assert [r.options["currency"] for r in db.scalars(select(JPReferenceEvent).order_by(JPReferenceEvent.revision))] == ["USD", "HKD"]
    assert db.scalar(select(func.count()).select_from(Source)) == 1  # Same bytes, separate interpretation evidence.


@pytest.mark.parametrize("value,expected", [("US$2.9 qty1000", "22.62"), ("28452(OPS)US$2.9100", "22.698"), ("2.9 USD", "22.62"), (0, "0")])
def test_confirmation_amount_not_an_unrelated_number(tmp_path, value, expected):
    from app.services.huakang_b_jp.prices import PriceTable
    path = tmp_path / "confirmed.xlsx"; path.write_bytes(confirmation(value, "USD"))
    obj = PriceTable(); obj.import_confirmed_xlsx(path)
    assert next(iter(obj.confirmed.values()))["price"] == expected


@pytest.mark.parametrize("amount,currency", [("HK$100=US$12.82", ""), ("price2.9 qty1000", "USD"), ("US$2.9", "HKD"), ("US$2.9.1", "USD"),
    ("US$2.90 or 3.50", "USD"), ("US$2.90 revised to 3.50", "USD"), ("US$2.90/3.50", "USD"),
    ("2.90 -> US$3.50", "USD")])
def test_ambiguous_confirmation_is_rejected(amount, currency):
    with pytest.raises(ValueError):
        intake._jp_validate_reference("confirmed", confirmation(amount, currency))


@pytest.mark.parametrize('kind', ['schedule', 'price', 'confirmed'])
@pytest.mark.parametrize('failure', ['none', 'load', 'iteration'])
def test_reference_import_releases_file_even_with_retained_traceback(tmp_path, monkeypatch, kind, failure):
    import builtins
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    from app.services.huakang_b_jp import workbooks
    from app.services.huakang_b_jp.prices import PriceTable
    from app.services.huakang_b_jp.schedule import ScheduleRef
    path = tmp_path / 'reference.xlsx'
    path.write_bytes(xlsx(kind))
    handles = []
    def tracked_open(*args, **kwargs):
        handle = builtins.open(*args, **kwargs)
        handles.append(handle)
        return handle
    monkeypatch.setattr(workbooks, 'open', tracked_open, raising=False)
    if failure == 'load':
        def fail_load(*args, **kwargs): raise ValueError('injected workbook load')
        monkeypatch.setattr(workbooks.openpyxl, 'load_workbook', fail_load)
    elif failure == 'iteration':
        original_iter = ReadOnlyWorksheet.iter_rows
        def fail_iter(self, *args, **kwargs):
            rows = original_iter(self, *args, **kwargs)
            yield next(rows)
            raise ValueError('injected row parsing')
        monkeypatch.setattr(ReadOnlyWorksheet, 'iter_rows', fail_iter)
    callback = {'schedule': ScheduleRef().load, 'price': PriceTable().import_xlsx,
                'confirmed': PriceTable().import_confirmed_xlsx}[kind]
    if failure != 'none':
        # Retain the exception/streaming generator references while deleting:
        # gc.collect() must not be needed for Windows temporary-file cleanup.
        with pytest.raises(ValueError, match='injected') as error:
            callback(path)
        assert error.traceback
    else:
        callback(path)
    assert handles and all(handle.closed for handle in handles)
    renamed = path.with_name('released.xlsx')
    path.rename(renamed)
    renamed.unlink()
