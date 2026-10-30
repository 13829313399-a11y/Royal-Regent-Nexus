"""Shipping-layout variants must retain field meaning, not merely column positions."""
from io import BytesIO
import openpyxl
import pytest

from app.services.indonesia_master_import import parse_material_book
from app.services.indonesia_master_workbook import named_piece_usage, price_currency
from app.schemas.indonesia_materials import MasterSheetMapping


def alternate_layout_book(change=None):
    book = openpyxl.Workbook(); sheet = book.active; sheet.title = '任意日期'
    sheet['A1'] = '出货资料'
    for col, value in {4: '货号\nItem No.', 5: '产品中文名称\nChinese name of the product',
        6: '产品英文名称\nEnglish name', 7: '规格\nspecification', 8: '类别\nkategori',
        9: '单位\nUnit', 10: '送货KG重量\nQuantity', 11: '送货数量\nProduct quantity', 12: '单位\nUnit',
        13: '采购单价（换成港币）\nPurchase price', 16: '净重总重Total Net weight（KGM）',
        19: '图片\nPicture', 20: '产品用途', 37: '供应商', 40: '采购单价', 41: '采购金额',
        43: '出口公司', 49: '备注'}.items():
        sheet.cell(2, col, value)
    for col, value in {4: 'SYNTH-P01', 5: '合成裁片-测试组件*2', 6: 'Synthetic component', 8: '合成类别',
        9: 'KGM', 10: 4.7, 11: 3600, 12: 'PCS', 16: 4.7, 40: 12.75, 41: 59.925}.items():
        sheet.cell(3, col, value)
    sheet['AN3'].number_format = '"US$"#,##0.00_);[Red]\\("US$"#,##0.00\\)'
    if change:
        change(sheet)
    output = BytesIO(); book.save(output)
    return output.getvalue()


def test_shifted_shipping_layout_recognizes_units_price_and_name_usage():
    result = parse_material_book(alternate_layout_book(), 'layout-49-column.xlsx')
    assert result['sheets'][0]['columns']['customs_unit'] == 9
    assert result['sheets'][0]['columns']['unit'] == 12
    assert result['sheets'][0]['columns']['price'] == 40
    row = result['rows'][0]
    assert row['source']['customs_unit'] == 'KGM'
    assert row['source']['unit'] == 'PCS'
    assert row['source']['price'] == '12.75'
    assert row['source']['currency'] == 'USD'
    assert row['evidence']['currency']['cell'] == 'AN3'
    assert row['source']['piece_usage'] == '2'
    assert row['source']['net_weight_kg'] == '0.00130556'
    assert row['evidence']['net_weight_kg']['inputs'] == [{'cell': 'P3', 'value': '4.7'}, {'cell': 'K3', 'value': '3600'}]
    assert row['source']['name'].endswith('组件*2')
    assert row['evidence']['piece_usage']['cell'] == 'E3'
    assert 'code' not in row['source'] and 'specification' not in row['source']


def test_inserted_columns_still_use_semantic_block_not_absolute_positions():
    result = parse_material_book(alternate_layout_book(lambda sheet: sheet.insert_cols(8)), '增加一列.xlsx')
    assert result['sheets'][0]['columns']['customs_unit'] == 10
    assert result['sheets'][0]['columns']['unit'] == 13
    assert result['rows'][0]['source']['price'] == '12.75'


def test_ambiguous_units_and_explicit_mapping_do_not_guess():
    result = parse_material_book(alternate_layout_book(lambda sheet: setattr(sheet['J2'], 'value', '未知数量')), '含糊.xlsx')
    assert 'customs_unit' not in result['rows'][0]['source']
    mapping = MasterSheetMapping(sheet='任意日期', header_row=2, columns={'item_no': 4, 'name': 5, 'price': 13})
    result = parse_material_book(alternate_layout_book(), '手工映射.xlsx', [mapping])
    assert 'customs_unit' not in result['rows'][0]['source']
    assert 'price' not in result['rows'][0]['source']
    assert 'piece_usage' not in result['rows'][0]['source']
    assert 'currency' not in result['rows'][0]['source']
    mapping.columns.update(customs_unit=9, unit=12)
    result = parse_material_book(alternate_layout_book(), '仅手工映射.xlsx', [mapping])
    assert 'net_weight_kg' not in result['rows'][0]['source']


@pytest.mark.parametrize('number_format, expected', [('"US$"#,##0.00', 'USD'), ('[$USD-409]0.00', 'USD'),
    ('"HK$"0.00', 'HKD'), ('"RMB"0.00', 'RMB'), ('"CNY"0.00', 'RMB'), ('"$"0.00', None),
    ('"¥"0.00', None), ('"USD"0.00;"HKD"0.00', None), ('0.00', None), ('"NOTUSD"0.00', None)])
def test_currency_never_comes_from_another_price_column_or_ambiguous_symbol(number_format, expected):
    assert price_currency(number_format) == expected


def test_currency_column_precedes_format_and_empty_hkd_price_is_not_a_default():
    def no_format(sheet):
        sheet['AN3'].number_format = '0.00'
    row = parse_material_book(alternate_layout_book(no_format), '未知币种.xlsx')['rows'][0]
    assert 'currency' not in row['source']
    def explicit(sheet):
        sheet['AU2'] = '币种'; sheet['AU3'] = 'HKD'
    row = parse_material_book(alternate_layout_book(explicit), '明确币种.xlsx')['rows'][0]
    assert row['source']['currency'] == 'HKD'


@pytest.mark.parametrize('name, expected', [('测试片*1', '1'), ('组件＊2', '2'), ('裁片*0.5', '0.5'), ('线*2', '2'),
    ('泡棉20*20*2', None), ('垫片20*2', None), ('标签*0', None), ('线*1000000000001', None), ('没有后缀', None)])
def test_name_suffix_is_not_a_dimension(name, expected):
    assert named_piece_usage(name) == expected


def test_shared_product_suffix_and_explicit_usage_stay_independent():
    result = parse_material_book(alternate_layout_book(lambda sheet: setattr(sheet['D3'], 'value', 'SYNTH-P01/SYNTH-P02')), '共用.xlsx')
    assert 'piece_usage' not in result['rows'][0]['source']
    assert '逐个货号确认' in result['rows'][0]['warnings'][0]
    def explicit(sheet):
        sheet['AU2'] = '单件用量'; sheet['AU3'] = 3
    result = parse_material_book(alternate_layout_book(explicit), '明确用量.xlsx')
    assert result['rows'][0]['source']['piece_usage'] == '3'


@pytest.mark.parametrize('cell, value', [('K3', 0), ('K3', None), ('K3', -1), ('K3', True), ('K3', '1e-40'), ('K3', '1e-100000000'),
    ('P3', '#VALUE!'), ('P3', '=J3'), ('P3', -1), ('L3', None)])
def test_invalid_weight_inputs_never_become_unit_weight(cell, value):
    result = parse_material_book(alternate_layout_book(lambda sheet: setattr(sheet[cell], 'value', value)), '待核对.xlsx')
    row = result['rows'][0]
    assert 'net_weight_kg' not in row['source']
    assert any('单个净重未换算' in warning for warning in row['warnings'])


def test_total_net_formula_uses_only_saved_result_and_records_both_coordinates():
    from test_indonesia_master_template import rewrite, ET, S
    def formula(sheet):
        sheet['P3'] = '=J3'
    def cache(parts):
        xml = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        node = next(node for node in xml.iter(S+'c') if node.get('r') == 'P3')
        node.find(S+'v').text = '4.7'
        parts['xl/worksheets/sheet1.xml'] = ET.tostring(xml)
    row = parse_material_book(rewrite(alternate_layout_book(formula), cache), '有缓存.xlsx')['rows'][0]
    assert row['source']['net_weight_kg'] == '0.00130556'
    assert row['evidence']['net_weight_kg']['inputs'][0] == {'cell': 'P3', 'value': '4.7', 'formula': '=J3', 'cached_value': '4.7'}
    def string_cache(parts):
        cache(parts)
        xml = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        next(node for node in xml.iter(S+'c') if node.get('r') == 'P3').set('t', 'str')
        parts['xl/worksheets/sheet1.xml'] = ET.tostring(xml)
    row = parse_material_book(rewrite(alternate_layout_book(formula), string_cache), '非数字缓存.xlsx')['rows'][0]
    assert 'net_weight_kg' not in row['source']


def test_shared_total_and_explicit_single_weight_cannot_be_replaced_by_a_derivation():
    def merged(sheet):
        sheet.merge_cells('P3:P4'); sheet['D4'] = 'SYNTH-P01'; sheet['E4'] = '另一物料'
    rows = parse_material_book(alternate_layout_book(merged), '共用重量.xlsx')['rows']
    assert all('net_weight_kg' not in row['source'] for row in rows)
    assert all(any('跨行合并' in warning for warning in row['warnings']) for row in rows)
    def explicit(sheet):
        sheet['AU2'] = '单个净重'; sheet['AU3'] = .25
    row = parse_material_book(alternate_layout_book(explicit), '明确单重.xlsx')['rows'][0]
    assert row['source']['net_weight_kg'] == '0.25'


@pytest.mark.parametrize('label, field, values', [('币种', 'currency', ('HKD', 'RMB')),
    ('单件用量', 'piece_usage', (3, 4)), ('单个净重', 'net_weight_kg', (.2, .3))])
def test_duplicate_explicit_columns_are_not_treated_as_absent(label, field, values):
    def duplicate(sheet):
        sheet['AU2'] = label; sheet['AV2'] = label
        sheet['AU3'] = values[0]; sheet['AV3'] = values[1]
    row = parse_material_book(alternate_layout_book(duplicate), '歧义表头.xlsx')['rows'][0]
    assert field not in row['source']
    assert any(label in warning and '多个同名列' in warning for warning in row['warnings'])
