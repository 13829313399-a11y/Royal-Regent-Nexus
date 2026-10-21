"""P4e contract tests using isolated ledgers and real XLSX readback."""
from io import BytesIO
from dataclasses import replace
import pytest
from openpyxl import load_workbook
from test_cutting_payments import (paid_setup, fees, price, fee_save, fee_decision, fee_accountant,
    client, user, command, BASE, flow, planned, production, ok, save, decision, reviewer, read)
from test_cutting_periods import (month_client, accounting, custody, bridge, actor, lock, act,
    draft, get as period_get, write)
from app.api import cutting_operations as api
from app.services import cutting_exports as exports
from app.services import cutting_export_workbook as workbooks
from test_cutting_periods import report
from test_cutting_materials import warehouse_post
from test_cutting_usage import read as usage_read
from zipfile import ZipFile
from hashlib import sha256
import json


def request(kind='payments', **kwargs):
    return dict(factory_id='huakang-c', kind=kind, month='2026-10') | kwargs


def download(p, body):
    preview = p.post(BASE+'/exports/preview', json=body)
    assert preview.status_code == 200, preview.text
    response = p.post(BASE+'/exports/download', json=body | {'token': preview.json()['token']})
    assert response.status_code == 200, response.text
    assert response.headers['cache-control'] == 'no-store'
    return load_workbook(BytesIO(response.content)), preview.json(), response


@pytest.fixture(autouse=True)
def fresh_auth(monkeypatch, client):
    # The base fixture replaces the dependency with synthetic identities; the final recheck uses the same identity.
    monkeypatch.setattr(api, 'get_current_user', lambda request, db: client.cutting_auth['user'])


def rows(wb, name):
    return list(wb[name].values)[2:]


def test_payment_export_exact_balances_statuses_and_trace(paid_setup):
    p=paid_setup
    one=save(p, evidence='=HYPERLINK("https://example.invalid","unsafe")')
    ok(decision(p,one,'submit'));reviewer(p);ok(decision(p,one,'confirm'))
    two=save(p, amount='2.50', allocations=[{'settlement_id':p.one,'amount':'2.50'}]);ok(decision(p,two,'submit'))
    wb, preview, response=download(p,request())
    assert rows(wb,'当前全期余额汇总') == [('CNY','40.00','5.00','35.00','2.50','32.50')]
    assert '不是历史月末余额' in dict(preview['metadata'])['口径']
    assert len(rows(wb,'所选月核销明细'))==2
    assert any('原冻结结算版本ID' in row for row in list(wb['所选月核销明细'].values))
    evidence=[cell for row in wb['所选月付款登记'] for cell in row if isinstance(cell.value,str) and cell.value.startswith('=HYPERLINK')]
    assert len(evidence)==1 and evidence[0].data_type=='s'
    assert all(cell.data_type!='f' for ws in wb for row in ws for cell in row)
    assert all(ws.print_title_rows=='$1:$2' for ws in wb)
    empty,_,_=download(p,request(month='2026-09'))
    assert rows(empty,'所选月付款登记')==[]
    assert rows(empty,'当前全期余额汇总')==rows(wb,'当前全期余额汇总')


def test_period_snapshot_stays_frozen_after_reprice_and_reclose(month_client):
    p=month_client;event=lock(p)
    body=request('period',event_id=event['id'])
    old, preview,_=download(p,body)
    act(p,'reopen')
    stale=p.post(BASE+'/exports/download',json=body|{'token':preview['token']})
    assert stale.status_code==409
    write(p,'price',issue_id=p.issue_id,unit_cost='4',currency='HKD',price_reference='反结调价')
    lock(p)
    historical, info,_=download(p,body)
    assert '历史失效' in dict(info['metadata'])['使用状态']
    for name in ('费用汇总','生产汇总','物料收发耗存','完工及在制归集','冻结来源索引'):
        assert rows(historical,name)==rows(old,name)
    current=period_get(p)
    new,_,_=download(p,request('period',event_id=current['event_id']))
    assert rows(new,'费用汇总')!=rows(old,'费用汇总')
    assert dict(info['metadata'])['快照校验值']==dict(preview['metadata'])['快照校验值']
    later=rows(historical,'后续失效与重结追溯')
    assert any(r[2]=='反结' and r[3] and r[4] and r[5] for r in later)
    assert any(r[2]=='复核锁月' for r in later)
    assert dict(info['metadata'])['当前有效冻结事件ID']==current['event_id']
    assert rows(old,'后续失效与重结追溯')==[]


def test_settlement_export_never_recalculates_quote_and_requires_confirmation(paid_setup):
    p=paid_setup
    value=ok(p.get(BASE+'/settlements/2026-10',params={'factory_id':'huakang-c','document_id':p.one}))
    body=request('settlement',document_id=p.one,event_id=value['event_id'])
    old,_,_=download(p,body)
    assert rows(old,'结算汇总')[0][5]=='20.00'
    ok(fee_decision(p,p.two,'reopen'));ok(fee_decision(p,p.one,'reopen'))
    historical, info,_=download(p,body)
    assert rows(historical,'验收计费明细')==rows(old,'验收计费明细')
    assert '历史失效' in dict(info['metadata'])['使用状态']
    assert any(r[2]=='反结' and r[5] for r in rows(historical,'后续失效与重结追溯'))
    assert dict(info['metadata'])['当前有效冻结事件ID']=='当前无有效冻结版本'
    assert p.post(BASE+'/exports/preview',json=body|{'event_id':'missing'}).status_code==404
    assert p.post(BASE+'/exports/preview',json=body|{'month':'2026-09'}).status_code==404
    first=value['history'][0]['id']
    assert p.post(BASE+'/exports/preview',json=body|{'event_id':first}).status_code==409


def test_export_permission_factory_bounds_and_draft_rejection(paid_setup):
    p=paid_setup
    for change in ({'factory_id':'huaxing'},{'month':'2026-13'},{'month':'2099-01'}, {'event_id':'injected'}):
        assert p.post(BASE+'/exports/preview',json=request()|change).status_code==422
    assert p.post(BASE+'/exports/download',json=request()).status_code==409
    for actions in [('read',),('read','cost_read'),('read','settlement_read')]:
        p.cutting_auth['user']=user(*actions,department='accounting')
        assert p.post(BASE+'/exports/preview',json=request()).status_code==403
    p.cutting_auth['user']=user('read','payment_read',department='production')
    assert p.post(BASE+'/exports/preview',json=request()).status_code==403


@pytest.mark.parametrize('change',['permission','session','balance'])
def test_generation_rechecks_identity_permissions_and_balance(paid_setup,monkeypatch,change):
    p=paid_setup;body=request();preview=ok(p.post(BASE+'/exports/preview',json=body))
    original=exports.package_isolated
    def altered(report, user_context):
        content=original(report,user_context)
        if change=='permission':p.cutting_auth['user']=user('read',department='accounting')
        elif change=='session':p.cutting_auth['user']=replace(p.cutting_auth['user'],id='another-user')
        else:save(p)
        return content
    monkeypatch.setattr(exports,'package_isolated',altered)
    response=p.post(BASE+'/exports/download',json=body|{'token':preview['token']})
    assert response.status_code==(409 if change=='balance' else 403)
    assert 'spreadsheetml' not in response.headers['content-type']


def test_workbook_limits_and_large_exact_values():
    report=dict(kind='payments',metadata=[],token='fixed',captured_at='2026-10-16',sheets=[exports.sheet('精度',['金额','货号'],[['9999999999999999.99','000001']])])
    wb=load_workbook(BytesIO(exports.render(report,user('read'))))
    assert rows(wb,'精度')==[('9999999999999999.99','000001')]
    report['sheets'][0]['rows']=[['x'*32768,'000001']]
    with pytest.raises(Exception) as too_long:exports.render(report,user('read'))
    assert too_long.value.status_code==413


def test_named_styles_keep_printing_and_text_contract():
    report = dict(kind='payments', metadata=[], token='fixed', captured_at='2026-10-16',
        sheets=[exports.sheet('样式', ['货号', '金额'], [['00001', '1.20'], ['=literal', '0.00']])])
    wb = load_workbook(BytesIO(exports.render(report, user('read'))))
    ws = wb['样式']
    assert rows(wb, '样式') == [('00001', '1.20'), ('=literal', '0.00')]
    assert ws['A4'].data_type == 's'
    for cell, bold, color in [(ws['A2'], True, '00FFFFFF'), (ws['A3'], False, '00173E43')]:
        assert cell.font.name == 'Microsoft YaHei' and cell.font.sz == 10
        assert cell.font.bold == bold and cell.font.color.rgb == color
        assert cell.alignment.horizontal == cell.alignment.vertical == 'center'
        assert cell.alignment.wrap_text
    assert ws['A2'].fill.fgColor.rgb == '00007F78'
    assert ws['A3'].fill.patternType is None
    assert ws['A4'].fill.fgColor.rgb == '00EFF7F7'
    assert ws.freeze_panes == 'B3' and ws.print_title_rows == '$1:$2'
    assert ws.column_dimensions['A'].width == 18 and ws.row_dimensions[1].height == 58
    assert ws.page_setup.orientation == 'landscape' and str(ws.page_setup.paperSize) == ws.PAPERSIZE_A3


def test_currency_and_explicit_zero_are_never_combined_or_lost():
    tables=exports.period_sheets(dict(costs={'HKD':{'period_normal':'0.00'},'CNY':{'period_normal':'10.01'}},
        orders=[],allocations=[],notices=[]))
    total=tables[0];index=total['headers'].index('本期正常费用')
    assert [(row[0],row[index]) for row in total['rows']]==[('CNY','10.01'),('HKD','0.00')]


def test_frozen_source_fields_periods_and_correction_links(month_client):
    p=month_client
    assert report(p,document_id=p.report_id,sets=8).status_code==200
    original=usage_read(p)['documents'][0]['id']
    write(p,'correct',document_id=original,issue_id=p.issue_id,consumed_quantity='2',loss_quantity='0.5',counted_remaining='1.5')
    returned=warehouse_post(p,'RETURN',original_id=p.issue_id,quantity='0.5',location_id='bin',business_date='2026-12-15')
    event=lock(p);snapshot=period_get(p)['data']['snapshot']
    wb,_,_=download(p,request('period',event_id=event['id']))
    trace={r[3]:r for r in rows(wb,'冻结来源索引')}
    production=[d for o in snapshot['orders'] for d in o['production_documents']]
    for document in production:
        for e in document['history']:
            row=trace[e['id']]
            assert row[5:8]==(e['data']['entry']['day'],e['actor_id'],e['created_at'])
            assert row[9]=='本期'
        posts=[e for e in document['history'] if e['action']=='post']
        assert len(posts)==2
        assert trace[posts[0]['id']][15]==posts[1]['id']
        assert trace[posts[1]['id']][14]==posts[0]['id']
        assert trace[posts[1]['id']][10]=='有效记录'
        assert trace[posts[0]['id']][10]=='历史／操作记录'
    assert trace[returned['id']][9]=='后续期间'
    assert trace[returned['id']][11]==p.issue_id
    assert any(r[12]==original for r in trace.values())
    assert any(r[2]==original and r[13] for r in trace.values())
    # Layout has separate compact summary and full-size, horizontally paginated details.
    assert wb['打印摘要'].page_setup.fitToWidth==1
    assert wb['完工及在制归集'].page_setup.scale==100
    assert not wb['完工及在制归集'].sheet_properties.pageSetUpPr.fitToPage
    assert wb['完工及在制归集'].print_title_cols=='$A:$A'
    assert wb['完工及在制归集'].col_breaks.count>0


def synthetic_report(count=20001):
    return dict(kind='payments',month='2026-10',metadata=[['口径','完整测试报表']],token='fixed',captured_at='2026-10-16',sheets=[
        exports.sheet('当前全期余额汇总',['币种','有效应付'],[['CNY','1.00']]),
        exports.sheet('明细',['流水号','数量'],[[f'{i:08d}','1.000000'] for i in range(count)])])


def test_automatic_volume_coverage_manifest_and_exact_text():
    report=synthetic_report()
    delivery=workbooks.delivery(report)
    assert delivery['format']=='zip' and delivery['volumes']>1
    data,extension=exports.package(report,user('read'))
    assert extension=='zip'
    collected=[];summaries=0
    with ZipFile(BytesIO(data)) as archive:
        manifest=json.loads(archive.read('manifest.json'))
        assert manifest['token']==report['token']
        for number, file in enumerate(manifest['files'],1):
            content=archive.read(file['name'])
            assert sha256(content).hexdigest()==file['sha256']
            wb=load_workbook(BytesIO(content))
            assert dict(rows(wb,'报表说明'))['分卷']==f"{number}/{delivery['volumes']}"
            assert len(rows(wb,'完整性清单'))==len(manifest['coverage'])
            summaries+='打印摘要' in wb.sheetnames
            for part in manifest['coverage']:
                if part['volume']!=number or part['source']!='明细':continue
                values=rows(wb,part['sheet'])
                assert len(values)==part['rows']<=workbooks.MAX_ROWS
                assert values==[tuple(r) for r in report['sheets'][1]['rows'][part['first']-1:part['last']]]
                collected.extend(values)
    assert collected==[tuple(r) for r in report['sheets'][1]['rows']]
    assert summaries==1


def test_export_limits_fail_whole_request_and_control_characters(monkeypatch):
    report=synthetic_report(2)
    monkeypatch.setattr(workbooks,'MAX_REPORT_CELLS',2)
    with pytest.raises(Exception) as too_many:workbooks.delivery(report)
    assert too_many.value.status_code==413 and '未生成部分' in too_many.value.detail
    monkeypatch.setattr(workbooks,'MAX_REPORT_CELLS',2000000)
    report['sheets'][1]['rows'][0][1]='bad\x01'
    with pytest.raises(Exception) as illegal:workbooks.delivery(report)
    assert illegal.value.status_code==422
    report['sheets'][1]['rows'][0][1]='1'
    monkeypatch.setattr(workbooks,'MAX_FILE_BYTES',1)
    with pytest.raises(Exception) as too_big:exports.package(report,user('read'))
    assert too_big.value.status_code==413 and '整份下载已取消' in too_big.value.detail


def test_wide_table_cell_budget_and_empty_table_coverage():
    report=synthetic_report(0)
    report['sheets'].append(exports.sheet('宽表',[f'字段{i}' for i in range(24)],[[0]*24]*20001))
    volumes=workbooks.plan(report);catalog=workbooks.manifest(report,volumes)
    assert len(volumes)>1
    for volume in volumes:
        assert sum((part['end']-part['start']+2)*len(report['sheets'][part['table']]['headers']) for part in volume)<workbooks.MAX_CELLS
    chunks=[r for r in catalog if r['source']=='宽表']
    assert chunks[0]['first']==1 and chunks[-1]['last']==20001
    assert all(b['first']==a['last']+1 for a,b in zip(chunks,chunks[1:]))
    assert sum(r['rows'] for r in chunks)==20001
    assert [r for r in catalog if r['source']=='明细'][0]['rows']==0


@pytest.mark.parametrize('change',['permission','balance'])
def test_volume_api_rechecks_before_returning_zip(paid_setup,monkeypatch,change):
    p=paid_setup;original_sheets=exports.payment_sheets;original_render=exports.render
    # Exercise reduced per-volume limits in this process; the subprocess uses
    # the same production constants and has its own contract/readback test.
    monkeypatch.setattr(exports, 'package_isolated', exports.package)
    def expanded(value,events):
        return original_sheets(value,events)+[exports.sheet('分卷测试',['序号','内容'],[[i,'长依据'*2000] for i in range(5)])]
    monkeypatch.setattr(exports,'payment_sheets',expanded)
    monkeypatch.setattr(workbooks,'MAX_TEXT_BYTES',300000)
    body=request();preview=ok(p.post(BASE+'/exports/preview',json=body))
    assert preview['delivery']['format']=='zip'
    response=p.post(BASE+'/exports/download',json=body|{'token':preview['token']})
    assert response.status_code==200 and response.headers['content-type']=='application/zip'
    assert response.headers['content-disposition'].endswith('.zip"')
    def changed(report,actor):
        result=original_render(report,actor)
        if change=='permission':p.cutting_auth['user']=user('read',department='accounting')
        else:save(p)
        return result
    monkeypatch.setattr(exports,'render',changed)
    response=p.post(BASE+'/exports/download',json=body|{'token':preview['token']})
    assert response.status_code==(403 if change=='permission' else 409)
    assert response.headers['content-type']=='application/json'
