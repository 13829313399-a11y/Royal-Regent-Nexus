from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile, ZIP_DEFLATED

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest
from pydantic import ValidationError

from app.schemas.spreadsheet_merge import MergeRequest
from app.services import spreadsheet_merge as service


def request(primary=None, secondary=None, **options):
    return MergeRequest.model_validate({
        'primary': {'name': '主.xlsx/数据', 'header_row': 3, 'headers': ['货号', '产品名称', '备注'], 'rows': primary or []},
        'secondary': {'name': '副.csv', 'headers': ['编码', '名称', '颜色'], 'rows': secondary or []},
        'mapping': [{'source': 0, 'target': 0}, {'source': 1, 'target': 1}],
        'key_columns': [0], **options})


def test_fill_conflict_resolution_and_audit_preserve_all_values():
    payload = request([['001', None, '旧'], ['002', '主值']], [['001', '补值', '红'], ['002', '副值', '蓝'], ['003', '新行', '绿']],
                      remarks_column=2, extra_columns_to_remarks=[2])
    result = service.merge(payload)
    assert result['rows'] == [['001', '补值', '旧'], ['002', '主值', '颜色：蓝'], ['003', '新行', '颜色：绿']]
    assert result['stats'] == {'primary_rows': 2, 'secondary_rows': 3, 'output_rows': 3, 'filled_cells': 2, 'added_rows': 1}
    assert result['provenance'][0] == [{'table': 'primary', 'name': '主.xlsx/数据', 'row': 4}, {'table': 'secondary', 'name': '副.csv', 'row': 2}]
    assert {c['id'] for c in result['conflicts']} == {'p0:s0:c2', 'p1:s1:c1'}
    with pytest.raises(HTTPException) as error:
        service.export(payload)
    assert error.value.status_code == 409
    resolved = payload.model_copy(update={'resolutions': {'p0:s0:c2': 'primary', 'p1:s1:c1': 'secondary'}})
    content, filename, _ = service.export(resolved)
    assert filename.endswith('ready.xlsx')
    book = load_workbook(BytesIO(content), data_only=False)
    assert book.sheetnames == ['导入数据', '说明与示例']
    assert book['导入数据']['A2'].value == '001'
    assert book['导入数据']['B3'].value == '副值'
    audit = str(list(book['说明与示例'].values))
    assert '红' in audit and '主值' in audit and '副值' in audit


def test_ambiguous_and_empty_keys_never_merge_or_disappear():
    payload = request([['A', '一'], ['A', '二'], [None, '空']], [['A', '副'], [None, '副空'], ['B', 'B1'], ['B', 'B2']])
    result = service.merge(payload)
    assert len(result['rows']) == 7
    assert not result['conflicts']
    assert len(result['issues']) == 7
    assert not result['ready_for_import']
    assert service.export(payload)[1].endswith('pending.xlsx')


def test_composite_keys_exact_and_no_keys_append():
    payload = request([['001', 'A']], [['001', 'B'], ['001', 'A']], key_columns=[0, 1])
    assert len(service.merge(payload)['rows']) == 2
    assert len(service.merge(payload.model_copy(update={'key_columns': []}))['rows']) == 3


def test_unmapped_columns_are_warned_and_in_audit():
    payload = request([], [['001', '产品', '丢失不可接受']])
    result = service.merge(payload)
    assert '颜色' in result['warnings'][0]
    book = load_workbook(BytesIO(service.export(payload)[0]))
    assert '丢失不可接受' in str(list(book['说明与示例'].values))


def test_source_review_column_blocks_ready_even_when_unmapped():
    payload = request([], [['001', '产品', '尺寸单位待核']])
    payload.secondary.headers[2] = '待核问题'
    result = service.merge(payload)
    assert not result['ready_for_import']
    assert result['issues'] == [{'code': 'SOURCE_REVIEW_REQUIRED', 'message': '副.csv 第 2 行：尺寸单位待核', 'rows': [0]}]
    assert service.export(payload)[1].endswith('pending.xlsx')


def test_audit_only_tables_retain_full_index_without_joining_merge():
    from app.schemas.spreadsheet_merge import MergeTable
    audit = MergeTable(name='原始.xlsx/原表文字索引', headers=['原表位置', '原始文字/完整备注'], rows=[['测试!3', '完整配卡原文' * 300]])
    payload = request([], [['A', '产品']], audit_tables=[audit.model_dump()])
    result = service.merge(payload)
    assert result['stats']['output_rows'] == 1 and not result['issues']
    book = load_workbook(BytesIO(service.export(payload)[0]))
    records = [r for r in book['说明与示例'].values if r[0] == '审计附表数据']
    assert records[0][4] == '完整配卡原文' * 300
    oversized = {'name': 'large', 'headers': [str(i) for i in range(50)], 'rows': [['x'] * 50 for _ in range(3000)]}
    with pytest.raises(ValidationError):
        request([], [], audit_tables=[oversized, oversized])
    excessive_rows = {'name': 'rows', 'headers': ['一列'], 'rows': [['x'] for _ in range(4000)]}
    with pytest.raises(ValidationError):
        request([], [], audit_tables=[excessive_rows] * 3)


def test_export_never_executes_user_formula_strings():
    payload = request([['=HYPERLINK("https://evil")', '=1+1', '@SUM(A1)']], [])
    book = load_workbook(BytesIO(service.export(payload)[0]), data_only=False)
    for sheet in book:
        for row in sheet:
            assert all(cell.data_type != 'f' for cell in row)
    assert book['导入数据']['A2'].value.startswith('=HYPERLINK')


def test_inspection_formula_cache_numeric_identifiers_and_style_only_tail():
    book = Workbook()
    sheet = book.active
    sheet.append(['货号', '值'])
    sheet.append([123, '=1+2'])
    sheet['A2'].number_format = '00000'
    sheet['XFD1048576'].number_format = '@'
    content = service._bytes(book)
    result = service.inspect(content, 'sample.xlsx')
    assert result['sheets'][0]['rows'] == [['货号', '值'], ['00123']]
    assert any('公式' in w and '1 个无有效缓存' in w for w in result['warnings'])
    assert any('数字货号' in w for w in result['warnings'])


def test_csv_keeps_zeroes_and_non_ascii():
    result = service.inspect('货号,名称\n00123,产品\n'.encode('utf-8-sig'), 'test.csv')
    assert result['sheets'][0]['rows'][1] == ['00123', '产品']


def test_zip_preflight_rejects_bombs_entities_duplicate_paths_and_out_of_bounds(monkeypatch):
    def archive(entries):
        output = BytesIO()
        with ZipFile(output, 'w', ZIP_DEFLATED) as z:
            for name, content in entries:
                z.writestr(name, content)
        return output.getvalue()
    for entries in [
        [('a.xml', '<!DOCTYPE x [<!ENTITY a "bomb">]><x/>')],
        [('../escape.xml', '<x/>')],
        [('same.xml', '<x/>'), ('same.xml', '<x/>')],
        [('xl/worksheets/sheet1.xml', f'<worksheet xmlns="{service.NS[1:-1]}"><sheetData><row><c r="IW1"><v>1</v></c></row></sheetData></worksheet>')],
    ]:
        with pytest.raises(HTTPException):
            service.preflight_xlsx(archive(entries))
    monkeypatch.setattr(service, 'MAX_EXPANDED', 20)
    with pytest.raises(HTTPException) as error:
        service.preflight_xlsx(archive([('data.xml', '<x>' + 'A' * 1000 + '</x>')]))
    assert error.value.status_code == 413


def test_sparse_rectangles_and_sheet_count_are_bounded():
    output = BytesIO()
    sparse = f'<worksheet xmlns="{service.NS[1:-1]}"><sheetData><row><c r="A1"><v>1</v></c></row><row><c r="IV5001"><v>1</v></c></row></sheetData></worksheet>'
    with ZipFile(output, 'w') as archive:
        archive.writestr('xl/worksheets/sheet1.xml', sparse)
    with pytest.raises(HTTPException) as error:
        service.preflight_xlsx(output.getvalue())
    assert error.value.status_code == 413
    output = BytesIO()
    with ZipFile(output, 'w') as archive:
        for i in range(51):
            archive.writestr(f'xl/worksheets/sheet{i}.xml', '<worksheet/>')
    with pytest.raises(HTTPException) as error:
        service.preflight_xlsx(output.getvalue())
    assert error.value.status_code == 413


def _rename_worksheet_part(content, new_path):
    output = BytesIO()
    with ZipFile(BytesIO(content)) as original, ZipFile(output, 'w', ZIP_DEFLATED) as changed:
        for member in original.infolist():
            data = original.read(member)
            if member.filename in {'xl/_rels/workbook.xml.rels', '[Content_Types].xml'}:
                data = data.replace(b'worksheets/sheet1.xml', new_path.removeprefix('xl/').encode())
            name = new_path if member.filename == 'xl/worksheets/sheet1.xml' else member.filename
            changed.writestr(name, data)
    return output.getvalue()


@pytest.mark.parametrize('part', ['xl/worksheets/data.xml', 'xl/worksheets/custom/data.xml'])
def test_custom_worksheet_part_names_preserve_values_and_adapter_candidates(part):
    from app.services.carton_packaging_adapter import adapt_carton_workbook
    book = Workbook()
    sheet = book.active
    sheet.title = '测试'
    sheet.append(['货号', '装箱数', '做箱SIZE inch'])
    sheet.append([None, None, '纸质', 'L', 'W', 'H'])
    sheet.append(['00123', 4, 'A=B', 12, 10, 8])
    content = _rename_worksheet_part(service._bytes(book), part)
    # Demonstrate the renamed part is a valid workbook relationship, not corrupt input.
    decoded = load_workbook(BytesIO(content), read_only=True)
    assert list(decoded.active.values)[2][0] == '00123'
    decoded.close()
    assert service.preflight_xlsx(content)[part] == (3, 6)
    inspected = service.inspect(content, 'custom-part.xlsx')
    assert inspected['sheets'][0]['rows'][2] == ['00123', 4, 'A=B', 12, 10, 8]
    adapted = adapt_carton_workbook(content)
    assert adapted['sheets'][0]['rows'][1][0] == '00123'
    assert adapted['sheets'][0]['rows'][1][10] == 4
    assert service.inspect(content, 'custom-part.xlsx', 'carton-packaging-source')['sheets'][0]['rows'][1][0] == '00123'


def test_custom_worksheet_parts_keep_rectangular_budget_and_unchecked_paths_fail_closed():
    from app.services.carton_packaging_adapter import adapt_carton_workbook
    book = Workbook()
    book.active['A1'] = '货号'
    book.active['IV5001'] = '远端数据'
    sparse = _rename_worksheet_part(service._bytes(book), 'xl/worksheets/data.xml')
    with pytest.raises(HTTPException) as error:
        service.inspect(sparse, 'sparse.xlsx')
    assert error.value.status_code == 413
    # Valid OOXML may place the worksheet elsewhere. Unsupported, unscanned
    # relationships must produce a clear error in both readers, never empty rows.
    book = Workbook()
    book.active.append(['货号', '数量'])
    book.active.append(['00123', 12])
    outside = _rename_worksheet_part(service._bytes(book), 'xl/custom/data.xml')
    decoded = load_workbook(BytesIO(outside), read_only=True)
    assert list(decoded.active.values)[1][0] == '00123'
    decoded.close()
    for read in (lambda: service.inspect(outside, 'outside.xlsx'), lambda: adapt_carton_workbook(outside)):
        with pytest.raises(HTTPException) as error:
            read()
        assert error.value.status_code == 422 and '未完成安全检查' in error.value.detail


def test_schema_rejects_unsafe_cells_and_invalid_mapping():
    with pytest.raises(ValidationError):
        request([['A', {'unexpected': 'object'}]], [])
    with pytest.raises(ValidationError):
        request([], [], mapping=[{'source': 0, 'target': 0}, {'source': 1, 'target': 0}])
    with pytest.raises(ValidationError):
        request([], [], extra_columns_to_remarks=[2])
    with pytest.raises(ValidationError):
        request([['A', float('nan')]], [])
    with pytest.raises(HTTPException):
        service.merge(request([], [], resolutions={'stale-id': 'primary'}))


def carton_request(rows, **options):
    from app.services.carton_master_import import HEADERS
    headers = HEADERS['configurations']
    return MergeRequest.model_validate({'primary': {'name': '系统模板', 'headers': headers, 'rows': rows},
        'secondary': {'name': '空表', 'headers': headers, 'rows': []}, 'profile': 'carton-packaging', **options})


def test_carton_profile_uses_real_parser_and_plan_without_business_database(monkeypatch):
    from app.services import carton_master_import as contract
    row = ['00123', '产品', '默认', '外箱', 'A33', 30, 20, 10, 'cm', '个', 24, '是', '启用', '', '2.1', '2.3']
    assert service.merge(carton_request([row]))['ready_for_import']
    missing = row[:]
    missing[1] = ''
    result = service.merge(carton_request([missing]))
    assert not result['ready_for_import']
    assert any('产品名称' in message for message in result['validation_errors'])
    assert service.export(carton_request([missing]))[1].endswith('pending.xlsx')
    numeric = row[:]
    numeric[0] = 123
    assert any('真实文本' in message for message in service.merge(carton_request([numeric]))['validation_errors'])
    duplicate = row[:]
    assert any('重复' in message for message in service.merge(carton_request([row, duplicate]))['validation_errors'])


def test_batch_export_max_1000_and_business_import_contract():
    from app.services import carton_master_import as contract
    rows = [[str(i).zfill(5), '产品', '默认', '外箱', 'A33', 30, 20, 10, 'cm', '个', 24, '否', '启用', '', None, None] for i in range(1001)]
    payload = carton_request(rows, batch_size=1000)
    content, filename, media = service.export(payload)
    assert filename.endswith('ready.zip') and media == 'application/zip'
    with ZipFile(BytesIO(content)) as archive:
        assert len(archive.namelist()) == 2
        counts = []
        for name in archive.namelist():
            entries, errors = contract.parse(archive.read(name), name, 'configurations')
            assert not errors
            counts.append(len(entries))
        assert counts == [1000, 1]


def test_carton_batching_keeps_noncontiguous_groups_atomic_and_provenance_correct():
    from app.services import carton_master_import as contract
    def row(code, paper):
        return [code, '产品', '组', paper, 'A33', 30, 20, 10, 'cm', '个', 24, '否', '启用', '', None, None]
    payload = carton_request([row('A', '外箱'), row('B', '外箱'), row('A', '内箱')], batch_size=2)
    content, filename, _ = service.export(payload)
    assert filename.endswith('ready.zip')
    with ZipFile(BytesIO(content)) as archive:
        first, second = [archive.read(name) for name in archive.namelist()]
    first_book = load_workbook(BytesIO(first))
    assert [r[0] for r in list(first_book['导入数据'].values)[1:]] == ['A', 'A']
    refs = [r for r in first_book['说明与示例'].values if r[0] == '输出行来源']
    assert [(r[1], r[4], r[6]) for r in refs] == [(2, 2, 2), (3, 4, 4)]
    entries, errors = contract.parse(first, 'first.xlsx', 'configurations')
    plan, _ = contract.plan(service._EmptyDatabase(), 'tool-validation', 'configurations', entries, errors)
    assert not plan['errors'] and plan['added'] == 1
    assert len(contract.parse(second, 'second.xlsx', 'configurations')[0]) == 1
    with pytest.raises(HTTPException) as error:
        service.export(payload.model_copy(update={'batch_size': 1}))
    assert '不能拆开' in error.value.detail
    blank_group, default_group = row('A', '外箱'), row('A', '内箱')
    blank_group[2], default_group[2] = '  ', '默认'
    with pytest.raises(HTTPException):
        service.export(carton_request([blank_group, default_group], batch_size=1))


def test_oversized_audit_export_is_pending_and_clearly_marked(monkeypatch):
    row = ['A', '产品', '组', '外箱', 'A33', 30, 20, 10, 'cm', '个', 24, '否', '启用', '', None, None]
    monkeypatch.setattr(service, 'BUSINESS_MAX_BYTES', 1)
    content, filename, _ = service.export(carton_request([row]))
    assert filename.endswith('pending.xlsx')
    book = load_workbook(BytesIO(content))
    assert '超过业务导入 5 MiB' in str(list(book['说明与示例'].values))


def test_authentication_feature_gate_and_no_service_database(monkeypatch):
    from app.api.spreadsheet_merge import router
    from app.core.config import settings
    from app.db import get_db
    from app.services.auth import get_current_user
    app = FastAPI()
    app.include_router(router)
    # Anonymous authentication fails before reading from this inert dependency.
    app.dependency_overrides[get_db] = lambda: SimpleNamespace()
    monkeypatch.setattr(settings, 'document_tools_enabled', True)
    client = TestClient(app)
    payload = request([], []).model_dump()
    for endpoint in ('preview', 'export'):
        assert client.post(f'/api/tools/spreadsheet-merge/{endpoint}', json=payload).status_code == 401
    assert client.post('/api/tools/spreadsheet-merge/inspect', files={'file': ('a.csv', b'a\n1')}).status_code == 401
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id='active-user')
    assert client.post('/api/tools/spreadsheet-merge/preview', json=payload).status_code == 200
    assert client.get('/api/tools/spreadsheet-merge/templates').status_code == 422
    assert client.get('/api/tools/spreadsheet-merge/templates/not-allowed').status_code == 422
    monkeypatch.setattr(settings, 'document_tools_enabled', False)
    assert client.post('/api/tools/spreadsheet-merge/preview', json=payload).status_code == 503
