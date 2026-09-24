"""Regression scenarios from independent review, against disposable databases."""
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from io import BytesIO
from uuid import uuid4
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from test_uv_ops import client,setup_task,confirm,post,read,payload,user_with
from test_uv_ops_edges import setup_agent,event
from app.models import uv_operations as m
from app.services.uv_operations import common as c, runtime, exports, imports


def day_shift(client,day):
    return post(client,'shifts',name='合成日期班次',business_date=day,start_at=day+'T08:00:00+08:00',end_at=day+'T16:00:00+08:00')


def test_backdated_rework_cannot_borrow_from_future_replenishment(client):
    data=setup_task(client,quantity=20);root=data['task'];first=data['shift']
    data['batch']=confirm(client,data,good=0,rework=10)['batch']
    data['task']=post(client,'rework-tasks',batch_id=data['batch']['id'],expected_version=data['batch']['version'],code='R1',quantity=10)
    data['shift']=day_shift(client,'2026-09-24')
    data['batch']=confirm(client,data,good=10,rework=0,source_bucket='rework')['batch']
    data['task']=root;data['shift']=day_shift(client,'2026-09-25')
    data['batch']=confirm(client,data,good=0,rework=10)['batch']
    child=post(client,'rework-tasks',batch_id=data['batch']['id'],expected_version=data['batch']['version'],code='R2',quantity=10)
    result=client.post('/api/uv-operations/production/confirm',json=payload(task_id=child['id'],batch_id=data['batch']['id'],shift_id=first['id'],expected_version=data['batch']['version'],pass_index=1,source_bucket='rework',processed=10,good=10,rework=0,scrap=0,pending=0,evidence='合成非法回填'))
    assert result.status_code==409 and result.json()['code']=='historical_balance'
    assert len(read(client,'production'))==3


def test_payroll_and_new_participant_are_atomically_ordered(client):
    data=setup_task(client,quantity=10);confirm(client,data,good=10,rework=0)
    fields=dict(task_id=data['task']['id'],shift_id=data['shift']['id'],start_at='2026-09-23T08:00:00+08:00',end_at='2026-09-23T09:00:00+08:00')
    post(client,'participations',**fields,employee_id='A',employee_name='合成 A')
    def writer(kind):
        path='wages/confirm' if kind==0 else 'participations'
        body=payload(task_id=data['task']['id'],shift_id=data['shift']['id']) if kind==0 else payload(**fields,employee_id='B',employee_name='合成 B')
        return client.post('/api/uv-operations/'+path,json=body)
    with ThreadPoolExecutor(2) as pool: results=list(pool.map(writer,[0,1]))
    assert results[0].status_code==200
    assert results[1].status_code in {200,409}
    allocations=read(client,'wage-allocations')
    assert len(allocations)==(2 if results[1].status_code==200 else 1)
    result=client.post('/api/uv-operations/participations',json=payload(**fields,employee_id='C',employee_name='合成 C'))
    assert result.status_code==409 and result.json()['code']=='payroll_frozen'


def test_split_merge_lineage_and_all_batches_must_finish(client):
    data=setup_task(client,quantity=120)
    split=post(client,'batches/'+data['batch']['id']+'/split',quantity=60,expected_version=data['batch']['version'],reason='合成两批分别现场流转')
    assert sum(x['quantity'] for x in read(client,'batches'))==120
    assert len(read(client,'batch-relations'))==2
    merged=post(client,'batches/'+split['batches'][0]['id']+'/merge',other_batch_id=split['batches'][1]['id'],other_version=split['batches'][1]['version'],expected_version=split['batches'][0]['version'],reason='合成两批尚未加工重新合并')
    batch=merged['batches'][0]
    split=post(client,'batches/'+batch['id']+'/split',quantity=60,expected_version=batch['version'],reason='合成再次拆分验证谱系')
    data['batch']=split['batches'][0]
    assert confirm(client,data,good=60,rework=0)['task']['status']=='in_progress'
    data['batch']=split['batches'][1]
    assert confirm(client,data,good=60,rework=0)['task']['status']=='completed'
    assert len(read(client,'batch-relations'))==6
    assert len(read(client,'tasks/'+data['task']['id'])['batch_relations'])==6
    assert sum(row['quantity'] for row in read(client,'batches'))==120


def test_cancelled_rework_cannot_be_resurrected(client):
    data=setup_task(client);data['batch']=confirm(client,data)['batch']
    child=post(client,'rework-tasks',batch_id=data['batch']['id'],expected_version=data['batch']['version'],code='CANCEL-ME',quantity=6)
    post(client,'rework-tasks/'+child['id']+'/cancel',expected_version=child['version'],reason='合成任务取消重开校验')
    result=client.post('/api/uv-operations/production/confirm',json=payload(task_id=child['id'],batch_id=data['batch']['id'],shift_id=data['shift']['id'],expected_version=data['batch']['version'],pass_index=1,source_bucket='rework',processed=6,good=6,rework=0,scrap=0,pending=0,evidence='合成取消后报产'))
    assert result.status_code==409 and result.json()['code']=='task_cancelled'


def test_handover_cannot_predate_production_and_return_cannot_predate_receipt(client):
    data=setup_task(client,quantity=10);batch=confirm(client,data,good=10,rework=0)['batch']
    fields=dict(batch_id=batch['id'],expected_version=batch['version'],quantity=10,target='合成下工序')
    before=client.post('/api/uv-operations/handovers',json=payload(**fields,business_date='2026-09-22'))
    assert before.status_code==409 and before.json()['code']=='historical_balance'
    handover=post(client,'handovers',**fields,business_date='2026-09-23')
    row=post(client,'handovers/'+handover['id']+'/receive',expected_version=handover['version'],quantity=10,business_date='2026-09-25',evidence='合成延后签收')['handover']
    result=client.post('/api/uv-operations/handovers/'+row['id']+'/return',json=payload(expected_version=row['version'],quantity=10,business_date='2026-09-24',evidence='合成不合法提前退回'))
    assert result.status_code==409 and result.json()['code']=='return_date'


def test_same_timestamp_old_sequence_and_rebinding_do_not_rewind_run(client):
    machine,agent,source,headers=setup_agent(client)
    later=event(machine,source,11,count='3',kind='job_complete')
    earlier=event(machine,source,10,count='1',kind='job_failed');earlier['observed_at']=later['observed_at']
    def upload(values):return client.post('/api/internal/uv-agent/events/batch',headers=headers,json=dict(schema_version=1,batch_id=uuid4().hex,events=values)).json()['results']
    assert upload([later,earlier])[1]['status']=='persisted'
    run=read(client,'runs')[0]
    assert run['state']=='completed' and Decimal(run['raw_count'])==3
    stream=event(machine,source,12)|dict(stream_id='restored-stream')
    rejected=upload([stream])[0];assert rejected['retryable'] and rejected['code']=='stream_rebind_required'
    bound=post(client,'sources',agent_id=agent['agent']['id'],machine_id=machine['id'],source_id=source['source_id'],adapter_type='simulator')
    newer=event(machine,bound,1,count='4',kind='job_complete')|dict(stream_id='restored-stream',observed_at='2026-09-23T08:01:00+08:00')
    assert upload([newer,stream])[0]['status']=='persisted'
    run=read(client,'runs')[0]
    assert run['last_binding_version']==bound['binding_version'] and Decimal(run['raw_count'])==4
    post(client,'sources/'+bound['id']+'/unbind',expected_version=bound['version'],reason='合成换绑释放来源校验')
    other=post(client,'machines',code='NEW-SOURCE',name='合成新机台')
    post(client,'sources',agent_id=agent['agent']['id'],machine_id=other['id'],source_id=source['source_id'],adapter_type='simulator')


def test_period_close_serializes_writer_and_unresolved_late_blocks_reclose(client):
    machine,agent,source,headers=setup_agent(client)
    with Session(client.uv_engine) as db:
        row=c.lock_period(db,'2026-09-01');version=row.version;db.commit()
    def writer(kind):
        return client.post('/api/uv-operations/'+('periods/2026-09/close' if kind==0 else 'expenses'),json=payload(expected_version=version,reason='合成月结并发') if kind==0 else payload(category='department',cost_amount='7',currency='CNY',business_date='2026-09-23',evidence='合成并发费用'))
    with ThreadPoolExecutor(2) as pool: responses=list(pool.map(writer,[0,1]))
    assert responses[0].status_code==200,responses[0].text
    assert responses[1].status_code in {200,409}
    closed=responses[0].json()['data']['period']
    client.post('/api/internal/uv-agent/events/batch',headers=headers,json=dict(schema_version=1,batch_id=uuid4().hex,events=[event(machine,source)]))
    opened=post(client,'periods/2026-09/reopen',expected_version=closed['version'],reason='合成月结重开检查迟到证据')
    result=client.post('/api/uv-operations/periods/2026-09/close',json=payload(expected_version=opened['version'],reason='合成重复封存'))
    assert result.status_code==409 and result.json()['code']=='late_evidence_unresolved'


def test_full_export_over_5000_records_and_formula_safe_ids(client):
    data=setup_task(client,quantity=6000)
    with Session(client.uv_engine) as db:
        original=db.get(m.UvOpsTask,data['task']['id'])
        values={column.name:getattr(original,column.name) for column in original.__table__.columns if column.name not in {'id','code','created_at','updated_at'}}
        values['quantity']=1
        db.add_all([m.UvOpsTask(id=f'synthetic-{i:06}',code='=DANGEROUS()' if i==0 else f'EXPORT-{i:06}',**values) for i in range(5001)])
        db.commit()
    job=post(client,'exports',kind='tasks',start_date='2026-09-01',end_date='2026-09-30')
    with Session(client.uv_engine) as db:
        row=db.get(m.UvOpsExportJob,job['id']);exports.build(db,client.uv_auth['user'],row)
        assert row.row_count==5002
        book=load_workbook(BytesIO(row.artifact),read_only=True)
        rows=list(book['UV记录'].rows);headers=[cell.value for cell in rows[0]]
        assert len(rows)==5003
        formula=next(row[headers.index('code')] for row in rows[1:] if row[headers.index('code')].value=='=DANGEROUS()')
        assert formula.data_type=='s'
        book.close()


def test_import_mapping_units_and_durable_large_preview(client):
    content='编号,名称,客户\nMAP-1,合成映射,\n'.encode()
    result=client.post('/api/uv-operations/imports/preview',data=dict(factory_id='huakang-a',template='products-v1',units_confirmed='true',field_mapping='{"code":"编号","name":"名称","customer":"客户"}'),files={'file':('mapping.csv',content,'text/csv')})
    assert result.status_code==200,result.text
    job=result.json()['data'];assert not job['errors']
    post(client,'imports/'+job['id']+'/commit',expected_version=job['version'])
    assert read(client,'products')[0]['code']=='MAP-1'
    content=('code,name\n'+''.join(f'BIG-{i},'+('合成'*40)+'\n' for i in range(1000))).encode()
    job=client.post('/api/uv-operations/imports/preview',data=dict(factory_id='huakang-a',template='products-v1',units_confirmed='true'),files={'file':('large.csv',content,'text/csv')}).json()['data']
    assert job['status']=='queued'
    with Session(client.uv_engine) as db:
        row=db.get(m.UvOpsImportJob,job['id']);imports.build_preview(db,client.uv_auth['user'],row);db.commit()
    result=read(client,'imports/'+job['id'])
    assert result['valid_count']==1000 and result['status']=='preview' and len(read(client,'products'))==1


def test_gang_run_cost_once_and_cross_midnight_time(client):
    first=setup_task(client,quantity=24);second=setup_task(client,quantity=12)
    machine,agent,source,headers=setup_agent(client)
    client.post('/api/internal/uv-agent/events/batch',headers=headers,json=dict(schema_version=1,batch_id=uuid4().hex,events=[event(machine,source,1,count='3',kind='job_complete')]))
    run=read(client,'runs')[0]
    # Both frozen layouts explicitly refer to the same physical 12-slot fixture.
    with Session(client.uv_engine) as db:
        row=db.get(m.UvOpsTask,second['task']['id']);row.process_snapshot=dict(row.process_snapshot,fixture_id=first['fixture']['id'])
        db.commit()
    matched=post(client,'runs/'+run['id']+'/match',expected_version=run['version'],evidence='合成拼版槽位人工确认',allocations=[dict(task_id=data['task']['id'],batch_id=data['batch']['id'],pass_index=1,slots=slots,full_boards=3,tail_pieces=0,share=share) for data,slots,share in [(first,list(range(1,9)),'0.666667'),(second,list(range(9,13)),'0.333333')]])
    assert [x['suggested_pieces'] for x in matched['allocations']]==[24,12]
    cost=post(client,'run-costs',run_id=run['id'],expected_version=matched['run']['version'],cost_amount='30',currency='CNY',business_date='2026-09-23',evidence='合成运行唯一直接成本池')
    assert sorted(Decimal(x['cost_amount']) for x in cost['cost_allocations'])==[Decimal('10'),Decimal('20')]
    duplicate=client.post('/api/uv-operations/run-costs',json=payload(run_id=run['id'],expected_version=matched['run']['version'],cost_amount='30',currency='CNY',business_date='2026-09-23',evidence='合成重复成本池'))
    assert duplicate.status_code==409
    with Session(client.uv_engine) as db:
        row=db.get(m.UvOpsRun,run['id']);row.started_at=c.ts('2026-09-22T22:00:00+08:00');row.ended_at=c.ts('2026-09-23T06:00:00+08:00');db.commit()
        result=runtime.report(db,'2026-09-22','2026-09-23')
    assert [Decimal(row['machine_seconds']) for row in result['by_day']]==[Decimal(7200),Decimal(21600)]
    assert sum(Decimal(row['allocated_machine_seconds']) for row in result['task_allocations'])==Decimal(28800)


def test_sse_rechecks_authorization_each_tick_and_projects_costs(client,monkeypatch):
    import asyncio,json
    from app.services.uv_operations import live
    setup_task(client)
    monkeypatch.setattr(live,'SessionLocal',lambda:Session(client.uv_engine))
    monkeypatch.setattr(live,'get_current_user',lambda request,db:client.uv_auth['user'])
    async def no_sleep(_):pass
    monkeypatch.setattr(live.asyncio,'sleep',no_sleep)
    class Request:
        async def is_disconnected(self):return False
    async def exercise():
        client.uv_auth['user']=user_with('read')
        stream=live.events(Request(),'huakang-a')
        initial=await anext(stream)
        assert 'event: reset' in initial and 'cost_price_snapshot' not in initial and 'payroll_policy_snapshot' not in initial
        client.uv_auth['user']=user_with()
        revoked=await anext(stream)
        assert 'event: access_revoked' in revoked
        await stream.aclose()
    asyncio.run(exercise())


def test_matched_batch_cannot_reshape_and_rework_match_uses_bound_batch(client):
    data=setup_task(client)
    machine,agent,source,headers=setup_agent(client)
    client.post('/api/internal/uv-agent/events/batch',headers=headers,json=dict(schema_version=1,batch_id=uuid4().hex,events=[event(machine,source,1)]))
    run=read(client,'runs')[0]
    allocation=dict(task_id=data['task']['id'],batch_id=data['batch']['id'],pass_index=1,slots=[1],full_boards=1,tail_pieces=0,share='1')
    post(client,'runs/'+run['id']+'/match',expected_version=run['version'],evidence='合成批次运行已确认',allocations=[allocation])
    result=client.post('/api/uv-operations/batches/'+data['batch']['id']+'/split',json=payload(expected_version=data['batch']['version'],quantity=60,reason='不应允许改变已确认运行来源'))
    assert result.status_code==409 and result.json()['code']=='batch_run_allocated'
    data['batch']=confirm(client,data)['batch']
    child=post(client,'rework-tasks',batch_id=data['batch']['id'],expected_version=data['batch']['version'],code='MATCH-REWORK',quantity=6)
    detail=read(client,'tasks/'+child['id'])
    assert [batch['id'] for batch in detail['batches']]==[data['batch']['id']]
    with Session(client.uv_engine) as db:
        second=c.add(db,m.UvOpsRun,client.uv_auth['user'],machine_id=machine['id'],binding_id=source['id'],native_identity=uuid4().hex,native_job_id='REWORK',started_at=m.now(),last_observed_at=m.now(),state='completed')
        db.commit(); second_id=second.id; version=second.version
    result=post(client,'runs/'+second_id+'/match',expected_version=version,evidence='合成返工运行归属确认',allocations=[dict(allocation,task_id=child['id'])])
    assert result['allocations'][0]['task_id']==child['id'] and result['production_created'] is False


def test_report_export_reuses_one_report_calculation(client,monkeypatch):
    from app.services.uv_operations import reports
    data=setup_task(client);confirm(client,data)
    job=post(client,'exports',kind='daily',start_date='2026-09-23',end_date='2026-09-23',selected_ids=[])
    original=reports.report;calls=[]
    def once(*args,**kwargs):
        calls.append(1)
        return original(*args,**kwargs)
    monkeypatch.setattr(reports,'report',once)
    with Session(client.uv_engine) as db:
        row=db.get(m.UvOpsExportJob,job['id']);exports.build(db,client.uv_auth['user'],row)
        book=load_workbook(BytesIO(row.artifact),read_only=True)
        assert {'UV记录','按币种核算','数据缺项','口径与范围'}<=set(book.sheetnames)
        book.close()
    assert len(calls)==1


def test_sunday_production_and_explicit_zero_cost_are_retained(client):
    data=setup_task(client,quantity=10)
    data['shift']=day_shift(client,'2026-09-27')
    confirm(client,data,good=10,rework=0)
    post(client,'expenses',category='direct',task_id=data['task']['id'],cost_amount='0',currency='CNY',business_date='2026-09-27',evidence='合成周日该项费用明确为零')
    result=client.get('/api/uv-operations/reports',params=dict(factory_id='huakang-a',start_date='2026-09-27',end_date='2026-09-27')).json()['data']
    assert result['totals']['good']==10
    assert Decimal(result['cost_by_currency']['CNY']['cost_amount'])==0
    assert not any(row['code']=='missing_cost_evidence' for row in result['cost_missing'])
