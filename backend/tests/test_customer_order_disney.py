from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.services.customer_order_huaxing import _disney_download_name
from app.services.huaxing_order_legacy import disney_schedule


def test_disney_parser_supports_theme_park_tdse_and_international_orders():
    theme_park = disney_schedule.parse_text(
        """
ORDER NUMBER ORIGINAL / CONFIRMATION PAGE
L-9743347 CHANGE 1
EARLIEST SHIP DATE LATEST SHIP DATE ORDER DATE SUPPLIER NUMBER
11/30/2026 12/7/2026 5/21/2026 40059618
REVISION NUMBER
1000128076 MM ASTRONAUT PULLBACK 2502 EA 1/EA 2.1400 401060937150 14.99
Multi INNER/PK: 6/EA
No Size CASE/PK: 6
""",
        "WDW_PO_L-9743347 (R1).pdf",
    )
    assert theme_park["revision"] == 1
    assert theme_park["rows"] == [{
        "po_number": "L-9743347",
        "customer_code": "DLR",
        "customer": "DLR（美国乐园）",
        "country": "美国乐园",
        "item": "1000128076",
        "description": "MM ASTRONAUT PULLBACK",
        "quantity": 2502,
        "case_pack": 6,
        "ship_date": "2026-11-30",
        "unit_price_usd": 2.14,
        "source_type": "theme_park",
    }]
    tdse = disney_schedule.parse_text(
        """
Disney Store Purchase Order
Original W5897
Order Ship Anticipate Cancel Date
29-APR- 03-OCT- 02-DEC-
09-OCT-2026
2026 2026 2026
Ship via Ocean
1 1000128076 COLOR N272 6 / 1
ADULT MICKEY 360 2.14
SIZE Ticketing
""",
        "PO#W5897 (R0).pdf",
    )
    assert tdse["rows"][0]["ship_date"] == "2026-10-03"
    assert tdse["rows"][0]["quantity"] == 360
    assert tdse["rows"][0]["case_pack"] == 6

    international = disney_schedule.parse_text(
        """
TYPE PAGE P.O. BILL TO:
Original 1 F00000000014945
ORDERED SHIP ON ANTICIPATE CANCEL AFTER
07/16/26 10/30/26 11/08/26 11/05/26
CLASSIFICATION
INTERNATIONALTOYWORLD: TWDISQ227141
STYLE DISNEY ITEM NUMBER UPC COST SIZE QUANTITY
1000128131 MONSTERS INC PULLBACK Multi 3.07 NO SIZE 402
Case Pack = 6, Inner Pack = 6
""",
        "F00000000014945 (R2).pdf",
    )
    assert international["revision"] == 2
    assert international["rows"][0]["customer"] == "菲律宾（INTERNATIONALTOYWORLD）"
    assert international["rows"][0]["unit_price_usd"] == 3.07


def test_disney_download_name_identifies_the_batch_and_generation_time():
    assert _disney_download_name(
        "2026-迪士尼排期 20260815_迪士尼新单.xlsx",
        "F00000000014373",
        generated_at=datetime(2026, 8, 17, 20, 5, 6),
    ) == (
        "2026-迪士尼排期 20260815_迪士尼新单_"
        "F00000000014373_20260817_200506.xlsx"
    )


def test_disney_merge_keeps_latest_revision_and_ignores_terms_attachment():
    old = disney_schedule.parse_text(
        """
Disney Store Purchase Order
Original W5898
Order Ship Anticipate Cancel Date
29-APR- 31-OCT- 02-JAN-
06-NOV-2026
2026 2026 2027
Ship via Ocean
1 1000128136 COLOR N272 6 / 1
FROZEN 804 2.23
SIZE Ticketing
""",
        "PO#W5898 (R0).pdf",
    )
    revised = disney_schedule.parse_text(
        old_text := """
Disney Store Purchase Order
Change 1 W5898
Order Ship Anticipate Cancel Date
29-APR- 03-NOV- 02-JAN-
06-NOV-2026
2026 2026 2027
Ship via Ocean
1 1000128136 COLOR N272 6 / 1
FROZEN 804 2.23
SIZE Ticketing
""",
        "PO#W5898 (R1).pdf",
    )
    terms = disney_schedule.parse_text(old_text, "PO#W5898 TL (R1).pdf")

    rows, warnings = disney_schedule.merge_revisions([old, revised, terms])

    assert len(rows) == 1
    assert rows[0]["ship_date"] == "2026-11-03"
    assert any("保留最新修订" in warning for warning in warnings)
    assert any("TL/条款附件" in warning for warning in warnings)


def _disney_template() -> bytes:
    workbook = Workbook()
    item = workbook.active
    item.title = disney_schedule.ITEM_SHEET
    headers = {
        2: "客出单日期", 3: "PO号", 4: "客名", 5: "產品編號",
        6: "产品名称", 7: "产品名称中文", 8: "规格", 9: "PO数量",
        10: "外箱装箱数", 13: "日期码", 14: "验货日期", 15: "走货期",
        24: "订单单价USD", 25: "单价HK$", 26: "出厂价", 27: "总金额USD",
        28: "总金额HK$", 29: "出厂价总金额HK$",
    }
    for column, title in headers.items():
        item.cell(3, column, title)
    for column in range(1, 30):
        item.cell(4, column)._style = item.cell(3, column)._style
        item.cell(4, column).font = Font(name="Arial", size=9)
        item.cell(4, column).fill = PatternFill("solid", fgColor="FFF2CC")
        item.cell(5, column).font = Font(name="Arial", size=9, bold=True)
    item["B4"] = datetime(2026, 1, 1)
    item["B4"].number_format = "yyyy-mm-dd"
    item["C4"] = "OLD-1"
    item["D4"] = "DLR（美国乐园）"
    item["E4"] = "100"
    item["F4"] = "OLD EN"
    item["G4"] = "旧产品"
    item["I4"] = 100
    item["J4"] = 10
    item["X4"] = 2
    item["Y4"] = "=X4*7.75"
    item["AA4"] = "=I4*X4"
    item["AB4"] = "=I4*Y4"
    item["AC4"] = "=Z4*I4"
    item["G5"] = "合计"
    item["I5"] = "=SUM(I4:I4)"
    item["E8"] = "合计"
    item["I8"] = "=SUM(I4:I7)"
    item["I2"] = "=SUM(I4:I7)"

    review = workbook.create_sheet(disney_schedule.REVIEW_SHEET)
    for column, title in enumerate(("", "客出单日期", "PO号", "客名", "產品編號", "數量", "装箱", "客要求走货期", "备注", "单价", "金额HKD"), 1):
        review.cell(1, column, title)
    for column in range(1, 12):
        review.cell(2, column).font = Font(name="Arial", size=8)
    for target_column, item_column in {2: "B", 3: "C", 4: "D", 5: "E", 6: "I", 7: "J", 8: "O", 10: "Y", 11: "AC"}.items():
        review.cell(2, target_column, f"={disney_schedule.ITEM_SHEET}!{item_column}4")
        review.cell(3, target_column, f"={disney_schedule.ITEM_SHEET}!{item_column}5")

    order = workbook.create_sheet(disney_schedule.ORDER_SHEET)
    for column in range(1, 48):
        order.cell(1, column, f"字段{column}")
        order.cell(2, column).font = Font(name="Arial", size=8)
    order["A2"] = f"={disney_schedule.ITEM_SHEET}!B4"
    order["C2"] = f"={disney_schedule.ITEM_SHEET}!C4"
    for target_column, item_column in {
        1: "B", 3: "C", 5: "C", 6: "D", 7: "H", 8: "E", 10: "G", 11: "I",
        38: "N", 39: "O", 40: "Y", 43: "J", 44: "M",
    }.items():
        order.cell(3, target_column, f"={disney_schedule.ITEM_SHEET}!{item_column}5")
    order["AO3"] = "=AN3*K3"
    order["J4"] = "接单合计"
    order["AO4"] = "=SUM(AO2:AO3)"
    order.print_area = "A1:AU4"

    preserved = workbook.create_sheet("历史资料")
    preserved["A1"] = "必须保留"
    preserved["B1"] = "=1+1"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_disney_export_groups_items_and_synchronizes_three_sheets(tmp_path: Path):
    source = _disney_template()
    output = tmp_path / "迪士尼新单.xlsx"
    records = [
        {
            "received_date": "2026-08-17", "po_number": "NEW-1",
            "customer": "WDW（美国乐园）", "country": "美国乐园",
            "item": "100", "description": "NEW EN", "product_name_zh": "旧产品",
            "quantity": 60, "case_pack": 10, "inspection_date": "2026-09-25",
            "ship_date": "2026-09-30", "unit_price_usd": 2.5,
            "date_code": "FAC-003843-26283", "factory_unit_price_hkd": 18,
        },
        {
            "received_date": "2026-08-17", "po_number": "NEW-2",
            "customer": "TDSE（欧洲）", "country": "欧洲",
            "item": "200", "description": "SECOND EN", "product_name_zh": "新产品",
            "quantity": 24, "case_pack": 6, "inspection_date": "2026-10-01",
            "ship_date": "2026-10-06", "unit_price_usd": 3,
            "date_code": "FAC-003843-26290", "factory_unit_price_hkd": 21,
        },
    ]

    disney_schedule.create_export(
        records,
        output,
        source,
        template_filename="迪士尼排期.xlsx",
    )

    rendered = load_workbook(output, data_only=False)
    assert rendered.sheetnames == ["ITEM表", "正单评审表", "接单表", "历史资料"]
    item = rendered["ITEM表"]
    assert item["C5"].value == "NEW-1"
    assert item["I6"].value == "=SUM(I4:I5)"
    assert item["C7"].value == "NEW-2"
    assert item["G8"].value == "合计"
    assert item["I8"].value == "=SUM(I7:I7)"
    assert item["E9"].value == "合计"
    assert item["I2"].value == "=SUM(I4:I8)"
    assert item["Y5"].value == "=X5*7.75"
    assert item["AC5"].value == "=Z5*I5"
    assert item["B5"].style_id == item["B4"].style_id

    review = rendered["正单评审表"]
    assert review["C3"].value == "=ITEM表!C5"
    assert review["C4"].value == "=ITEM表!C6"
    assert review["C5"].value == "=ITEM表!C7"
    assert review["C6"].value == "=ITEM表!C8"
    order = rendered["接单表"]
    assert order["C3"].value == "=ITEM表!C5"
    assert order["F3"].value == "=ITEM表!D5"
    assert order["G3"].value == "=ITEM表!H5"
    assert order["C4"].value == "=ITEM表!C6"
    assert order["J7"].value == "接单合计"
    assert rendered["历史资料"]["A1"].value == "必须保留"
    assert rendered["历史资料"]["B1"].value == "=1+1"
    rendered.close()


def test_disney_existing_contract_without_subtotal_is_extended_directly(tmp_path: Path):
    workbook = load_workbook(BytesIO(_disney_template()), data_only=False)
    item = workbook[disney_schedule.ITEM_SHEET]
    item["G5"] = None
    item["I5"] = None
    review = workbook[disney_schedule.REVIEW_SHEET]
    for column in range(1, 12):
        review.cell(3, column).value = None
    order = workbook[disney_schedule.ORDER_SHEET]
    for column in range(1, 48):
        order.cell(3, column).value = None
    order["J3"] = "接单合计"
    order["AO3"] = "=SUM(AO2:AO2)"
    order["J4"] = None
    order["AO4"] = None
    direct_source = BytesIO()
    workbook.save(direct_source)
    workbook.close()

    output = tmp_path / "迪士尼现有合同续行.xlsx"
    disney_schedule.create_export(
        [{
            "received_date": "2026-08-17",
            "po_number": "OLD-1",
            "customer": "DLR（美国乐园）",
            "country": "美国乐园",
            "item": "100",
            "description": "OLD EN",
            "product_name_zh": "旧产品",
            "quantity": 20,
            "case_pack": 10,
            "inspection_date": "2026-12-20",
            "ship_date": "2026-12-25",
            "unit_price_usd": 2,
            "date_code": "FAC-003843-26350",
            "factory_unit_price_hkd": 15,
        }],
        output,
        direct_source.getvalue(),
        template_filename="迪士尼排期.xlsx",
    )

    rendered = load_workbook(output, data_only=False)
    item = rendered[disney_schedule.ITEM_SHEET]
    assert item["C4"].value == "OLD-1"
    assert item["C5"].value == "OLD-1"
    assert item["C6"].value is None
    assert item["G6"].value is None
    assert item["E9"].value == "合计"
    assert item["B5"].style_id == item["B4"].style_id
    assert item["AC5"].value == "=Z5*I5"
    assert rendered[disney_schedule.REVIEW_SHEET]["C3"].value == "=ITEM表!C5"
    assert rendered[disney_schedule.ORDER_SHEET]["C3"].value == "=ITEM表!C5"
    assert rendered[disney_schedule.ORDER_SHEET]["J4"].value == "接单合计"
    rendered.close()
