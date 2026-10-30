from dataclasses import replace
from datetime import datetime, timedelta
from io import BytesIO
import importlib.util
from pathlib import Path
from urllib.parse import parse_qs, urlsplit

from alembic.migration import MigrationContext
from alembic.operations import Operations
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pydantic import SecretStr
import pytest
from sqlalchemy import create_engine, select, text
from sqlalchemy.exc import DatabaseError
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.api.online_documents import router, authenticated_user
from app.core.config import settings
from app.db import get_db
from app.models.auth import (AuthUser, AuthSession, AuthIamState, EmployeeProfile, AuthPermission,
    AuthPermissionMetadata, AuthRole, AuthRoleMetadata, AuthRolePermission, AuthUserRole,
    AuthRoleBindingMetadata, AuthUserPermissionOverride)
from app.models.identity import IamOrgUnit, IamOrgDepartment, EmployeeAssignment, IamRoleVersion
from app.models.online_documents import OnlineDocument as Document, OnlineDocumentRevision as Revision
from app.services import auth, online_documents as service, online_document_protocol as protocol
from app.services.document_tools import storage

BASE = "/api/tools/online-documents"
SCOPE = "?factory_id=huaxing"


def context(user="owner", factory="huaxing", epoch=1):
    return auth.AuthContext(id=user, username=user, display_name=user, roles=(), role_codes=(), permissions=frozenset(),
                            factory_scopes=(factory,), department_scopes=("engineering",),
                            profile=auth.AuthProfileContext(primary_factory_id=factory, primary_department="engineering", confirmation_status="confirmed"),
                            identity={"employment_epoch": epoch})


def workbook(formula="=SUM(2,3)"):
    book = Workbook(); book.active["A1"] = formula
    stream = BytesIO(); book.save(stream)
    return stream.getvalue()


@pytest.fixture
def api(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "online_documents_enabled", True)
    monkeypatch.setattr(settings, "onlyoffice_jwt_secret", SecretStr("test-only-secret-" * 4))
    monkeypatch.setattr(settings, "onlyoffice_public_url", "http://office.example/office")
    monkeypatch.setattr(settings, "onlyoffice_internal_url", "http://onlyoffice")
    monkeypatch.setattr(settings, "document_tools_storage_dir", str(tmp_path / "assets"))
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    for model in (AuthUser, AuthSession, AuthIamState, EmployeeProfile, AuthPermission,
                  AuthPermissionMetadata, AuthRole, AuthRoleMetadata, AuthRolePermission,
                  AuthUserRole, AuthRoleBindingMetadata, AuthUserPermissionOverride,
                  IamOrgUnit, IamOrgDepartment, EmployeeAssignment, IamRoleVersion):
        model.__table__.create(engine)
    path = Path(__file__).parents[1] / "alembic/versions/20261026_0177_online_documents.py"
    spec = importlib.util.spec_from_file_location("online_migration", path)
    migration = importlib.util.module_from_spec(spec); spec.loader.exec_module(migration)
    with engine.begin() as conn:
        migration.op = Operations(MigrationContext.configure(conn)); migration.upgrade()
        path = Path(__file__).parents[1] / "alembic/versions/20261026_0179_online_document_management.py"
        spec = importlib.util.spec_from_file_location("management_migration", path)
        management_migration = importlib.util.module_from_spec(spec); spec.loader.exec_module(management_migration)
        management_migration.op = Operations(MigrationContext.configure(conn)); management_migration.upgrade()
        spec = importlib.util.spec_from_file_location("folder_migration", Path(__file__).parents[1] / "alembic/versions/20261026_0181_online_document_folders.py")
        folder_migration = importlib.util.module_from_spec(spec); spec.loader.exec_module(folder_migration)
        folder_migration.op = Operations(MigrationContext.configure(conn)); folder_migration.upgrade()
        spec = importlib.util.spec_from_file_location("folder_version_migration", Path(__file__).parents[1] / "alembic/versions/20261026_0183_document_folder_version.py")
        folder_version_migration = importlib.util.module_from_spec(spec); spec.loader.exec_module(folder_version_migration)
        folder_version_migration.op = Operations(MigrationContext.configure(conn)); folder_version_migration.upgrade()
    contexts = {who: context(who) for who in ("owner", "editor", "viewer", "stranger")}
    with Session(engine) as db:
        db.add(AuthIamState(key="identity_mutation_lock", value_json="{}", updated_at="2026-10-26T00:00:00Z"))
        for who in contexts:
            db.add(AuthUser(id=who, username=who, display_name=who, password_salt="x", password_hash="x", status="active"))
            db.add(EmployeeProfile(user_id=who, primary_factory_id="huaxing", primary_department="engineering", confirmation_status="confirmed"))
            db.add(AuthSession(id=who, user_id=who, token_hash=auth.hash_session_token(who), status="active", expires_at=(datetime.now()+timedelta(days=1)).strftime("%Y-%m-%d %H:%M:%S")))
        db.commit()
    monkeypatch.setattr(auth, "build_auth_context", lambda db, user: contexts[user.id])
    monkeypatch.setattr(service.collaborative_sheets, "_eligible_accounts", lambda db, factory: [{"id": key} for key, user in contexts.items() if user.profile.primary_factory_id == factory])
    app = FastAPI(); app.include_router(router)
    current = {"user": "owner"}
    app.dependency_overrides[authenticated_user] = lambda: contexts[current["user"]]
    def database():
        with Session(engine) as db:
            yield db
    app.dependency_overrides[get_db] = database
    with TestClient(app) as client:
        client.cookies.set("rr_session", "owner")
        yield client, engine, contexts, current, monkeypatch
    engine.dispose()


def put_grants(client, url, *, json):
    # Send the actual owner's latest ACL version unless a test supplies a stale one.
    detail = client.get(url.replace('/grants', '')).json()
    payload = {**json, 'grants': [{'epoch': 1, **g} for g in json['grants']]}
    payload.setdefault('version', detail.get('share_version', ''))
    return client.put(url, json=payload)


def new(client):
    response = client.post(BASE + SCOPE, json={"title": "公式核验", "format": "xlsx"})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def open_editor(client, document):
    response = client.post(f"{BASE}/{document}/editor{SCOPE}")
    assert response.status_code == 200, response.text
    return response.json()["config"]


def callback(client, document, payload, **extra):
    return client.post(f"{BASE}/{document}/callback", json={**payload, **extra, "token": protocol.sign(payload)})


def test_private_factory_epoch_and_view_permissions(api):
    client, engine, users, current, monkey = api
    doc = new(client)
    current["user"] = "stranger"
    assert client.get(BASE + SCOPE).json()["items"] == []
    assert client.get(f"{BASE}/{doc}/download{SCOPE}").status_code == 404
    current["user"] = "owner"
    grants = [{"user_id": "viewer", "access": "view"}, {"user_id": "editor", "access": "edit"}]
    assert put_grants(client, f"{BASE}/{doc}/grants{SCOPE}", json={"grants": grants}).status_code == 200
    current["user"] = "viewer"; client.cookies.set("rr_session", "viewer")
    config = open_editor(client, doc)
    assert config["document"]["permissions"]["edit"] is False
    assert client.post(f"{BASE}/{doc}/save{SCOPE}").status_code == 403
    assert put_grants(client, f"{BASE}/{doc}/grants{SCOPE}", json={"grants": []}).status_code == 403
    with Session(engine) as db:
        assert not db.get(Document, doc).editor_key  # Viewing never locks sharing.
    users["viewer"] = context("viewer", epoch=2)
    assert client.get(f"{BASE}/{doc}{SCOPE}").status_code == 404
    users["viewer"] = context("viewer", factory="huakang-a")
    assert client.get(f"{BASE}/{doc}{SCOPE}").status_code == 403
    current["user"] = "owner"; users["owner"] = context("owner", epoch=2)
    assert client.get(f"{BASE}/{doc}{SCOPE}").status_code == 404


def test_durable_save_source_preservation_order_and_session_fence(api):
    client, engine, users, current, monkey = api
    doc = new(client); original = client.get(f"{BASE}/{doc}/download{SCOPE}").content
    config = open_editor(client, doc); key = config["document"]["key"]
    token = parse_qs(urlsplit(config["document"]["url"]).query)["token"][0]
    assert client.get(f"{BASE}/{doc}/source", params={"token": token}).content == original
    assert put_grants(client, f"{BASE}/{doc}/grants{SCOPE}", json={"grants": []}).status_code == 409
    monkey.setattr(protocol, "command", lambda payload: {"error": 0})
    save = client.post(f"{BASE}/{doc}/save{SCOPE}").json()["request_id"]
    data = workbook("=SUM(7,8)")
    monkey.setattr(protocol, "fetch_saved", lambda url: data)
    payload = {"key": key, "status": 6, "forcesavetype": 0, "userdata": save, "users": ["owner_1"], "url": "http://onlyoffice/cache/files/test.xlsx"}
    assert callback(client, doc, payload).json() == {"error": 0}
    assert client.get(f"{BASE}/{doc}{SCOPE}").json()["saved_request"] == save
    assert load_workbook(BytesIO(client.get(f"{BASE}/{doc}/download{SCOPE}").content)).active["A1"].value == "=SUM(7,8)"
    assert client.get(f"{BASE}/{doc}/source", params={"token": token}).content == original
    assert open_editor(client, doc)["document"]["key"] == key
    # Replayed old force-save is ignored even if its file differs.
    monkey.setattr(protocol, "fetch_saved", lambda url: workbook("=0"))
    assert callback(client, doc, payload).json() == {"error": 0}
    assert client.get(f"{BASE}/{doc}{SCOPE}").json()["revision"] == 2
    assert callback(client, doc, {**payload, "status": 2}).json() == {"error": 0}
    new_key = open_editor(client, doc)["document"]["key"]
    assert new_key != key
    assert callback(client, doc, {**payload, "status": 2}).json() == {"error": 0}
    with Session(engine) as db:
        versions = db.scalars(select(Revision).where(Revision.document_id == doc).order_by(Revision.revision)).all()
        assert len(versions) == 3
        assert storage.resolve(versions[0].storage_key).read_bytes() == original
        assert db.get(Document, doc).editor_key == new_key
        with pytest.raises(DatabaseError):
            db.execute(text("UPDATE online_document_revisions SET actor_id='tampered'"))


@pytest.mark.parametrize("boundary", ["epoch", "disabled", "logout", "factory", "view"])
def test_callback_rechecks_identity_and_session(api, boundary):
    client, engine, users, current, monkey = api
    doc = new(client); key = open_editor(client, doc)["document"]["key"]
    if boundary == "epoch": users["owner"] = context("owner", epoch=2)
    if boundary == "factory": users["owner"] = context("owner", factory="huakang-a")
    with Session(engine) as db:
        if boundary == "disabled": db.get(AuthUser, "owner").status = "disabled"
        if boundary == "logout": db.get(AuthSession, "owner").status = "revoked"
        if boundary == "view":
            record = db.get(Document, doc); record.participants = {"owner_1": {**record.participants["owner_1"], "edit": False}}
        db.commit()
    monkey.setattr(protocol, "fetch_saved", lambda url: pytest.fail("Must reject before fetching"))
    result = callback(client, doc, {"key": key, "status": 2, "users": ["owner_1"], "url": "http://onlyoffice/cache/files/x"})
    assert result.status_code == 403
    with Session(engine) as db: assert db.get(Document, doc).revision == 1


def test_signed_payload_not_unsigned_body_and_source_purpose(api):
    client, engine, users, current, monkey = api
    doc = new(client); config = open_editor(client, doc)
    assert client.post(f"{BASE}/{doc}/callback", json={"status": 2}).status_code == 403
    monkey.setattr(protocol, "fetch_saved", lambda url: pytest.fail("Unsigned fields must never cause a save"))
    assert callback(client, doc, {"key": config["document"]["key"], "status": 1}, status=2, url="http://onlyoffice/cache/files/x").json() == {"error": 0}
    assert client.get(f"{BASE}/{doc}/source", params={"token": config["token"]}).status_code == 403
    assert client.post(f"{BASE}/{doc}/callback", content=b"x" * 140000).status_code == 413


@pytest.mark.parametrize("url", ["http://evil.example/cache/files/a", "http://onlyoffice/admin", "http://onlyoffice@evil/cache/files/a", "http://onlyoffice/cache/files/%252e%252e/admin", "http://onlyoffice/cache/files/a#bad", "http://onlyoffice:81/cache/files/a", None])
def test_save_url_rejects_arbitrary_destinations(api, url):
    with pytest.raises(HTTPException): protocol.cache_url(url)


def test_proxy_cache_url_and_token_expiry(api):
    assert protocol.cache_url("http://office.example/office/cache/files/a?md5=abc") == "http://onlyoffice/cache/files/a?md5=abc"
    with pytest.raises(HTTPException): protocol.verify(protocol.sign({"exp": 1}))
    token = protocol.sign({"key": "test"})
    with pytest.raises(HTTPException): protocol.verify(token[:-2] + "ZZ")


def test_upload_and_new_word_and_unsupported_input(api):
    client, *_ = api
    assert client.post(BASE + SCOPE, json={"format": "docx"}).status_code == 201
    data = workbook()
    response = client.post(BASE + "/upload", data={"factory_id": "huaxing"}, files={"file": ("upload.xlsx", data)})
    assert response.status_code == 201
    assert client.get(f"{BASE}/{response.json()['id']}/download{SCOPE}").content == data
    assert client.post(BASE + "/upload", data={"factory_id": "huaxing"}, files={"file": ("bad.xlsx", b"not-a-workbook")}).status_code == 422


def test_uncertain_force_save_never_becomes_false_success(api):
    client, engine, users, current, monkey = api
    doc = new(client); key = open_editor(client, doc)["document"]["key"]
    commands = []
    def command(payload):
        commands.append(payload)
        return {"error": -1 if len(commands) == 1 else 4}
    monkey.setattr(protocol, "command", command)
    assert client.post(f"{BASE}/{doc}/save{SCOPE}").status_code == 502
    request = commands[0]["userdata"]
    assert client.post(f"{BASE}/{doc}/save{SCOPE}").status_code == 502
    assert commands[1]["userdata"] == request
    detail = client.get(f"{BASE}/{doc}{SCOPE}").json()
    assert detail["pending_save"] == request and detail["saved_request"] != request
    monkey.setattr(protocol, "fetch_saved", lambda url: workbook("=7+8"))
    assert callback(client, doc, {"key": key, "status": 6, "forcesavetype": 0, "userdata": request,
                                  "users": ["owner_1"], "url": "http://onlyoffice/cache/files/x"}).json() == {"error": 0}
    assert client.get(f"{BASE}/{doc}{SCOPE}").json()["saved_request"] == request
    assert client.post(f"{BASE}/{doc}/save{SCOPE}").status_code == 200


def test_stale_force_save_failure_does_not_poison_new_request(api):
    client, engine, users, current, monkey = api
    doc = new(client); key = open_editor(client, doc)["document"]["key"]
    monkey.setattr(protocol, "command", lambda payload: {"error": 0})
    request = client.post(f"{BASE}/{doc}/save{SCOPE}").json()["request_id"]
    assert callback(client, doc, {"key": key, "status": 7, "forcesavetype": 0, "userdata": "old"}).json() == {"error": 0}
    detail = client.get(f"{BASE}/{doc}{SCOPE}").json()
    assert detail["pending_save"] == request and not detail["save_error"]


def test_release_unopened_checks_info_without_forcing_save(api):
    client, engine, users, current, monkey = api
    doc = new(client); key = open_editor(client, doc)["document"]["key"]
    calls = []
    def command(payload):
        calls.append(payload)
        return {"error": 0 if len(calls) == 1 else 1}
    monkey.setattr(protocol, "command", command)
    assert client.post(f"{BASE}/{doc}/release-unopened{SCOPE}").json() == {"released": False}
    assert client.post(f"{BASE}/{doc}/release-unopened{SCOPE}").json() == {"released": True}
    assert calls == [{"c": "info", "key": key}] * 2


@pytest.mark.parametrize('factory', ['group', 'huaxing'])
def test_prepare_sharing_releases_confirmed_orphan_before_revoke(api, factory):
    client, engine, users, current, monkey = api
    suffix = '?factory_id=' + factory
    doc = client.post(BASE + suffix, json={'format': 'xlsx'}).json()['id']
    path = f'{BASE}/{doc}'
    shared = put_grants(client, path + '/grants' + suffix, json={'grants': [{'user_id': 'viewer', 'access': 'view'}]}).json()
    client.post(path + '/editor' + suffix)
    before = client.get(path + '/download' + suffix).content
    calls = []
    monkey.setattr(protocol, 'command', lambda p: calls.append(p) or {'error': 1})
    payload = {'grants': [], 'version': shared['share_version']}
    assert client.post(path + '/prepare-sharing' + suffix, json=payload).json() == {'ready': True}
    detail = client.get(path + suffix).json()
    assert not detail['editing'] and detail['grants'] == shared['grants']
    assert client.get(path + '/download' + suffix).content == before
    assert len(calls) == 1 and calls[0]['c'] == 'info'
    assert client.put(path + '/grants' + suffix, json=payload).status_code == 200
    current['user'] = 'viewer'
    assert client.get(path + '/download' + suffix).status_code == 404


def test_prepare_sharing_retains_last_editor_authority_until_final_save(api):
    client, engine, users, current, monkey = api
    doc = new(client); path = f'{BASE}/{doc}'
    shared = put_grants(client, path + '/grants' + SCOPE, json={'grants': [{'user_id': 'editor', 'access': 'edit'}]}).json()
    current['user'] = 'editor'; client.cookies.set('rr_session', 'editor')
    key = open_editor(client, doc)['document']['key']
    current['user'] = 'owner'; client.cookies.set('rr_session', 'owner')
    calls = []
    monkey.setattr(protocol, 'command', lambda p: calls.append(p) or {'error': 0, 'users': []})
    payload = {'grants': [], 'version': shared['share_version']}
    assert client.post(path + '/prepare-sharing' + SCOPE, json=payload).json() == {'ready': False}
    assert calls[:2] == [{'c': 'info', 'key': key}, {'c': 'drop', 'key': key, 'users': ['editor_1']}]
    assert calls[2]['c'] == 'forcesave' and calls[2]['key'] == key
    assert client.put(path + '/grants' + SCOPE, json=payload).status_code == 409
    assert client.get(path + SCOPE).json()['grants'] == shared['grants']
    saved = workbook('=SUM(100,200)')
    monkey.setattr(protocol, 'fetch_saved', lambda url: saved)
    final = {'key': key, 'status': 2, 'users': ['editor_1'], 'url': 'http://onlyoffice/cache/files/final'}
    assert callback(client, doc, final).json() == {'error': 0}
    assert client.get(path + '/download' + SCOPE).content == saved
    assert client.put(path + '/grants' + SCOPE, json=payload).status_code == 200
    current['user'] = 'editor'; client.cookies.set('rr_session', 'editor')
    assert client.post(path + '/editor' + SCOPE).status_code == 404
    assert callback(client, doc, final).json() == {'error': 0}


@pytest.mark.parametrize('status', [2, 4])
def test_prepare_sharing_handles_final_callback_before_drop_response(api, status):
    client, engine, users, current, monkey = api
    doc = new(client); path = f'{BASE}/{doc}'
    put_grants(client, path + '/grants' + SCOPE, json={'grants': [{'user_id': 'viewer', 'access': 'view'}]})
    key = open_editor(client, doc)['document']['key']
    version = client.get(path + SCOPE).json()['share_version']
    saved = workbook('=SUM(9,9)')
    monkey.setattr(protocol, 'fetch_saved', lambda url: saved)
    def command(payload):
        if payload['c'] == 'drop':
            assert callback(client, doc, {'key': key, 'status': status, 'users': ['owner_1'], 'url': 'http://onlyoffice/cache/files/final'}).json() == {'error': 0}
        return {'error': 0}
    monkey.setattr(protocol, 'command', command)
    assert client.post(path + '/prepare-sharing' + SCOPE, json={'grants': [], 'version': version}).json() == {'ready': True}
    assert not client.get(path + SCOPE).json()['editing']
    if status == 2:
        assert client.get(path + '/download' + SCOPE).content == saved


@pytest.mark.parametrize('failure', ['info-timeout', 'info-error', 'drop-timeout', 'drop-error', 'pending', 'save-error'])
def test_prepare_sharing_never_discards_uncertain_content(api, failure):
    client, engine, users, current, monkey = api
    doc = new(client); path = f'{BASE}/{doc}'
    original = put_grants(client, path + '/grants' + SCOPE, json={'grants': [{'user_id': 'viewer', 'access': 'view'}]}).json()['grants']
    key = open_editor(client, doc)['document']['key']
    version = client.get(path + SCOPE).json()['share_version']
    def command(payload):
        stage = payload['c']
        if failure == stage + '-timeout':
            raise TimeoutError()
        if failure == stage + '-error':
            return {'error': 2}
        if failure in {'pending', 'save-error'}:
            # Save state can change while info is in flight, after the first check.
            with Session(engine) as db:
                record = db.get(Document, doc)
                if failure == 'pending': record.pending_save = 'unacknowledged'
                else: record.save_error = 'unconfirmed'
                db.commit()
            return {'error': 1}
        return {'error': 0}
    monkey.setattr(protocol, 'command', command)
    response = client.post(path + '/prepare-sharing' + SCOPE, json={'grants': [], 'version': version})
    assert response.status_code == (409 if failure in {'pending', 'save-error'} else 502)
    with Session(engine) as db:
        record = db.get(Document, doc)
        assert record.editor_key == key and record.participants
        assert record.grants == original
        if failure == 'pending': assert record.pending_save == 'unacknowledged'
        if failure == 'save-error': assert record.save_error == 'unconfirmed'


@pytest.mark.parametrize('boundary, expected', [('epoch', 403), ('session', 401), ('account', 401), ('acl', 409), ('key', 200), ('recipient', 409)])
def test_prepare_sharing_rechecks_owner_and_session_before_drop(api, boundary, expected):
    client, engine, users, current, monkey = api
    doc = new(client); path = f'{BASE}/{doc}'
    key = open_editor(client, doc)['document']['key']
    version = client.get(path + SCOPE).json()['share_version']
    calls = []
    def command(payload):
        calls.append(payload)
        if boundary == 'epoch': users['owner'] = context('owner', epoch=2)
        if boundary == 'recipient': users['viewer'] = context('viewer', epoch=2)
        with Session(engine) as db:
            if boundary == 'session': db.get(AuthSession, 'owner').status = 'revoked'
            if boundary == 'account': db.get(AuthUser, 'owner').status = 'disabled'
            if boundary == 'acl': db.get(Document, doc).grants = [{'user_id': 'viewer', 'name': 'viewer', 'epoch': 1, 'access': 'view'}]
            if boundary == 'key': db.get(Document, doc).editor_key = 'replacement-key'
            db.commit()
        return {'error': 0}
    monkey.setattr(protocol, 'command', command)
    intended = [{'user_id': 'viewer', 'epoch': 1, 'access': 'view'}] if boundary == 'recipient' else []
    response = client.post(path + '/prepare-sharing' + SCOPE, json={'grants': intended, 'version': version})
    assert response.status_code == expected, response.text
    assert calls == [{'c': 'info', 'key': key}]
    if boundary == 'key': assert response.json() == {'ready': False}


def test_prepare_sharing_validates_recipient_and_acl_before_disconnecting(api):
    client, engine, users, current, monkey = api
    doc = new(client); path = f'{BASE}/{doc}'
    open_editor(client, doc)
    version = client.get(path + SCOPE).json()['share_version']
    calls = []
    monkey.setattr(protocol, 'command', lambda p: calls.append(p) or {'error': 0})
    assert client.post(path + '/prepare-sharing' + SCOPE, json={'grants': [], 'version': 'old'}).status_code == 409
    assert client.post(path + '/prepare-sharing' + SCOPE, json={'grants': [{'user_id': 'viewer', 'epoch': 2, 'access': 'view'}], 'version': version}).status_code == 409
    current['user'] = 'stranger'; client.cookies.set('rr_session', 'stranger')
    assert client.post(path + '/prepare-sharing' + SCOPE, json={'grants': [], 'version': version}).status_code == 404
    assert calls == []


def test_shared_editor_cannot_end_connections_for_sharing(api):
    client, engine, users, current, monkey = api
    doc = new(client); path = f'{BASE}/{doc}'
    shared = put_grants(client, path + '/grants' + SCOPE, json={'grants': [{'user_id': 'editor', 'access': 'edit'}]}).json()
    current['user'] = 'editor'; client.cookies.set('rr_session', 'editor')
    open_editor(client, doc)
    calls = []
    monkey.setattr(protocol, 'command', lambda p: calls.append(p) or {'error': 0})
    assert client.post(path + '/prepare-sharing' + SCOPE, json={'grants': [], 'version': shared['share_version']}).status_code == 403
    assert calls == []


def test_startup_requires_explicit_migration_even_for_a_new_store(api):
    from app import db as database
    client, migrated, users, current, monkey = api
    empty = create_engine("sqlite://")
    monkey.setattr(database, "engine", empty)
    with pytest.raises(RuntimeError, match="20261026_0183"):
        database.ensure_online_documents_schema_ready()
    monkey.setattr(database, "engine", migrated)
    database.ensure_online_documents_schema_ready()
    empty.dispose()


def test_group_workspace_without_employee_profile_is_private_by_default(api):
    client, engine, users, current, monkey = api
    users['owner'] = replace(context(), profile=None, factory_scopes=(), department_scopes=(), permissions=frozenset({'system:users'}))
    personal = '?factory_id=group'
    response = client.post(BASE + personal, json={'format': 'xlsx', 'title': '个人文档'})
    assert response.status_code == 201
    doc = response.json()['id']
    assert client.get(BASE + personal).json()['items'][0]['id'] == doc
    assert {p['id'] for p in client.get(BASE + '/recipients' + personal).json()['users']} == {'editor', 'viewer', 'stranger'}
    assert put_grants(client, f'{BASE}/{doc}/grants{personal}', json={'grants': []}).status_code == 200
    assert client.get(BASE + SCOPE).status_code == 403
    assert client.get(BASE + '?factory_id=unknown').status_code == 422
    config = client.post(f'{BASE}/{doc}/editor{personal}').json()['config']
    token = parse_qs(urlsplit(config['document']['url']).query)['token'][0]
    assert client.get(f'{BASE}/{doc}/source', params={'token': token}).status_code == 200
    monkey.setattr(protocol, 'command', lambda payload: {'error': 0})
    nonce = client.post(f'{BASE}/{doc}/save{personal}').json()['request_id']
    saved = workbook('=SUM(7,8)')
    monkey.setattr(protocol, 'fetch_saved', lambda url: saved)
    assert callback(client, doc, {'key': config['document']['key'], 'status': 6, 'forcesavetype': 0,
        'userdata': nonce, 'users': ['owner_1'], 'url': 'http://onlyoffice/cache/files/personal'}).json() == {'error': 0}
    assert client.get(f'{BASE}/{doc}/download{personal}').content == saved
    # Group browsing and even another internal account confer no implicit access.
    current['user'] = 'stranger'
    assert client.get(BASE + personal).json()['items'] == []
    for action in ('', '/download'):
        assert client.get(f'{BASE}/{doc}{action}{personal}').status_code == 404
    assert client.post(f'{BASE}/{doc}/save{personal}').status_code == 404
    current['user'] = 'owner'
    users['owner'] = replace(users['owner'], identity={'employment_epoch': 2})
    assert client.get(f'{BASE}/{doc}/download{personal}').status_code == 404
    assert client.get(f'{BASE}/{doc}/source', params={'token': token}).status_code in (403, 404)


def test_unavailable_account_cannot_use_personal_workspace(api):
    client, engine, users, current, monkey = api
    users['owner'] = replace(context(), account_available=False)
    for path in ('', '/recipients'):
        assert client.get(BASE + path + '?factory_id=group').status_code == 403
    assert client.post(BASE + '?factory_id=group', json={'format': 'docx'}).status_code == 403


@pytest.mark.parametrize('factory', ['group', 'huaxing'])
def test_rename_preserves_file_content_versions_and_sharing(api, factory):
    client, engine, users, current, monkey = api
    suffix = '?factory_id=' + factory
    doc = client.post(BASE + suffix, json={'format': 'xlsx', 'title': '改名前'}).json()['id']
    path = f'{BASE}/{doc}'
    shared = put_grants(client, path + '/grants' + suffix, json={'grants': [{'user_id': 'viewer', 'access': 'view'}]}).json()
    original = client.get(path + '/download' + suffix).content
    with Session(engine) as db:
        original_key = db.get(Document, doc).storage_key
    response = client.patch(path + '/title' + suffix, json={'title': '  十月成品汇总.xlsx  ', 'expected_title': '改名前', 'expected_version': shared['state_version']})
    assert response.status_code == 200, response.text
    renamed = response.json()
    assert renamed['title'] == '十月成品汇总' and renamed['format'] == 'xlsx'
    assert renamed['grants'] == shared['grants'] and renamed['share_version'] == shared['share_version']
    assert renamed['revision'] == shared['revision'] == 1
    downloaded = client.get(path + '/download' + suffix)
    assert downloaded.content == original
    from urllib.parse import unquote
    assert '十月成品汇总.xlsx' in unquote(downloaded.headers['content-disposition'])
    with Session(engine) as db:
        assert db.get(Document, doc).storage_key == original_key
        assert len(db.scalars(select(Revision).where(Revision.document_id == doc)).all()) == 1
    current['user'] = 'viewer'
    assert client.get(path + suffix).json()['title'] == '十月成品汇总'


@pytest.mark.parametrize('name', ['', '  ', '.xlsx', 'a/b', 'a\nb'])
def test_rename_rejects_invalid_names_without_changing_title(api, name):
    client, engine, users, current, monkey = api
    doc = new(client)
    path = f'{BASE}/{doc}'
    assert client.patch(path + '/title' + SCOPE, json={'title': name, 'expected_title': '公式核验', 'expected_version': 1}).status_code == 422
    assert client.get(path + SCOPE).json()['title'] == '公式核验'


def test_rename_requires_owner_and_rejects_stale_updates(api):
    client, engine, users, current, monkey = api
    doc = new(client)
    path = f'{BASE}/{doc}'
    assert put_grants(client, path + '/grants' + SCOPE, json={'grants': [{'user_id': 'viewer', 'access': 'view'}, {'user_id': 'editor', 'access': 'edit'}]}).status_code == 200
    payload = {'title': '新名称', 'expected_title': '公式核验', 'expected_version': client.get(path + SCOPE).json()['state_version']}
    for who, status in [('viewer', 403), ('editor', 403), ('stranger', 404)]:
        current['user'] = who
        assert client.patch(path + '/title' + SCOPE, json=payload).status_code == status
    current['user'] = 'owner'
    assert client.patch(path + '/title' + SCOPE, json={**payload, 'expected_title': '旧窗口名称'}).status_code == 409
    open_editor(client, doc)
    commands = []
    monkey.setattr(protocol, 'command', lambda payload: commands.append(payload) or {'error': 0})
    with Session(engine) as db:
        record = db.get(Document, doc)
        before = (record.editor_key, record.editor_source, record.participants, record.storage_key, record.revision)
    assert client.patch(path + '/title' + SCOPE, json=payload).status_code == 200
    assert commands == [{'c': 'meta', 'key': before[0], 'meta': {'title': '新名称.xlsx'}}]
    with Session(engine) as db:
        record = db.get(Document, doc)
        assert (record.editor_key, record.editor_source, record.participants, record.storage_key, record.revision) == before
    assert client.get(path + SCOPE).json()['title'] == '新名称'
    users['owner'] = context('owner', epoch=2)
    assert client.patch(path + '/title' + SCOPE, json=payload).status_code == 404


def test_group_sharing_read_edit_save_downgrade_and_revoke(api):
    client, engine, users, current, monkey = api
    group = '?factory_id=group'
    users['owner'] = replace(users['owner'], profile=None, permissions=frozenset({'system:users'}))
    users['editor'] = context('editor', factory='huakang-a')
    doc = client.post(BASE + group, json={'format': 'xlsx'}).json()['id']
    path = f'{BASE}/{doc}'
    grants = [{'user_id': 'viewer', 'access': 'view'}, {'user_id': 'editor', 'access': 'edit'}]
    assert put_grants(client, path + '/grants' + group, json={'grants': grants}).status_code == 200
    current['user'] = 'viewer'; client.cookies.set('rr_session', 'viewer')
    assert client.get(BASE + group).json()['items'][0]['access'] == 'view'
    view = client.post(path + '/editor' + group).json()['config']
    assert view['editorConfig']['mode'] == 'view'
    assert view['document']['permissions']['edit'] is False
    view_token = parse_qs(urlsplit(view['document']['url']).query)['token'][0]
    assert client.get(path + '/download' + group).status_code == 200
    assert client.post(path + '/save' + group).status_code == 403
    assert put_grants(client, path + '/grants' + group, json={'grants': []}).status_code == 403
    assert 'grants' not in client.get(path + group).json()
    current['user'] = 'editor'; client.cookies.set('rr_session', 'editor')
    edit = client.post(path + '/editor' + group).json()['config']
    assert edit['document']['permissions']['edit'] is True
    assert put_grants(client, path + '/grants' + group, json={'grants': []}).status_code == 403
    current['user'] = 'owner'
    assert put_grants(client, path + '/grants' + group, json={'grants': []}).status_code == 409
    current['user'] = 'editor'
    monkey.setattr(protocol, 'command', lambda p: {'error': 0})
    nonce = client.post(path + '/save' + group).json()['request_id']
    saved = workbook('=SUM(11,22)')
    monkey.setattr(protocol, 'fetch_saved', lambda url: saved)
    payload = {'key': edit['document']['key'], 'status': 6, 'forcesavetype': 0,
               'userdata': nonce, 'users': ['editor_1'], 'url': 'http://onlyoffice/cache/files/shared'}
    assert callback(client, doc, payload).json() == {'error': 0}
    assert client.get(path + '/download' + group).content == saved
    assert callback(client, doc, {**payload, 'status': 2}).json() == {'error': 0}
    current['user'] = 'owner'
    grants[1]['access'] = 'view'
    assert put_grants(client, path + '/grants' + group, json={'grants': grants}).status_code == 200
    current['user'] = 'editor'
    assert client.post(path + '/editor' + group).json()['config']['editorConfig']['mode'] == 'view'
    assert client.post(path + '/save' + group).status_code == 403
    current['user'] = 'viewer'; client.cookies.set('rr_session', 'viewer')
    view = client.post(path + '/editor' + group).json()['config']
    view_token = parse_qs(urlsplit(view['document']['url']).query)['token'][0]
    assert client.get(path + '/source', params={'token': view_token}).status_code == 200
    current['user'] = 'owner'
    assert put_grants(client, path + '/grants' + group, json={'grants': []}).status_code == 200
    for who in ('viewer', 'editor', 'stranger'):
        current['user'] = who
        client.cookies.set('rr_session', who)
        assert client.get(BASE + group).json()['items'] == []
        assert client.get(path + '/download' + group).status_code == 404
        assert client.post(path + '/editor' + group).status_code == 404
    assert client.get(path + '/source', params={'token': view_token}).status_code == 404


@pytest.mark.parametrize('boundary', ['epoch', 'disabled', 'left', 'supplier', 'no_internal_assignment', 'unconfirmed'])
def test_group_grants_recheck_current_membership_on_reads_and_callbacks(api, boundary):
    client, engine, users, current, monkey = api
    group = '?factory_id=group'
    doc = client.post(BASE + group, json={'format': 'xlsx'}).json()['id']
    path = f'{BASE}/{doc}'
    assert put_grants(client, path + '/grants' + group, json={'grants': [{'user_id': 'editor', 'access': 'edit'}]}).status_code == 200
    current['user'] = 'editor'; client.cookies.set('rr_session', 'editor')
    config = client.post(path + '/editor' + group).json()['config']
    token = parse_qs(urlsplit(config['document']['url']).query)['token'][0]
    if boundary == 'epoch': users['editor'] = context('editor', epoch=2)
    if boundary == 'disabled': users['editor'] = replace(users['editor'], account_available=False)
    if boundary == 'left': users['editor'] = replace(users['editor'], identity={'employment_epoch': 1, 'employment_status': 'left'})
    if boundary == 'supplier': users['editor'] = replace(users['editor'], permissions=frozenset({'carton_supplier:read'}))
    if boundary == 'no_internal_assignment': users['editor'] = replace(users['editor'], identity={'employment_epoch': 1, 'identity_mode': 'v2'})
    if boundary == 'unconfirmed': users['editor'] = replace(users['editor'], profile=replace(users['editor'].profile, confirmation_status='needs_review'))
    assert client.get(path + group).status_code in (403, 404)
    assert client.get(path + '/source', params={'token': token}).status_code in (403, 404)
    assert client.post(path + '/save' + group).status_code in (403, 404)
    monkey.setattr(protocol, 'fetch_saved', lambda url: pytest.fail('Unauthorized editor must not fetch saved bytes'))
    assert callback(client, doc, {'key': config['document']['key'], 'status': 2, 'users': ['editor_1'], 'url': 'http://onlyoffice/cache/files/x'}).status_code == 403
    with Session(engine) as db: assert db.get(Document, doc).revision == 1


def test_group_recipients_use_internal_roster_and_recheck_target(api):
    client, engine, users, current, monkey = api
    group = '?factory_id=group'
    with Session(engine) as db:
        db.add(IamOrgUnit(id='huaxing', name='华星', kind='factory', legacy_factory_id='huaxing'))
        db.get(AuthUser, 'stranger').status = 'disabled'
        db.get(EmployeeProfile, 'viewer').employment_status = 'left'
        db.commit()
    response = client.get(BASE + '/recipients' + group)
    assert response.status_code == 200, response.text
    assert response.json()['users'] == [{'id': 'editor', 'display_name': 'editor', 'epoch': 1, 'department': '工程部', 'organization': '华星'}]
    doc = client.post(BASE + group, json={'format': 'xlsx'}).json()['id']
    for target in ('owner', 'stranger', 'viewer', 'unknown'):
        assert put_grants(client, f'{BASE}/{doc}/grants{group}', json={'grants': [{'user_id': target, 'access': 'view'}]}).status_code == 422
    # A target changing after roster discovery is checked again, not trusted from the picker.
    users['editor'] = replace(users['editor'], permissions=frozenset({'carton_supplier:read'}))
    assert put_grants(client, f'{BASE}/{doc}/grants{group}', json={'grants': [{'user_id': 'editor', 'access': 'edit'}]}).status_code == 403
    current['user'] = 'editor'
    assert client.get(BASE + '/recipients' + group).status_code == 403
    own = client.post(BASE + group, json={'format': 'xlsx'}).json()['id']
    assert put_grants(client, f'{BASE}/{own}/grants{group}', json={'grants': []}).status_code == 403


def test_group_does_not_relax_concrete_factory_sharing(api):
    client, engine, users, current, monkey = api
    users['editor'] = context('editor', factory='huakang-a')
    doc = new(client)
    assert put_grants(client, f'{BASE}/{doc}/grants{SCOPE}', json={'grants': [{'user_id': 'editor', 'access': 'edit'}]}).status_code == 422
    current['user'] = 'editor'
    assert client.get(f'{BASE}/{doc}{SCOPE}').status_code == 403
    assert client.get(f'{BASE}/{doc}?factory_id=group').status_code == 404


def test_legacy_unconfirmed_empty_identity_cannot_receive_or_manage_group_sharing(api):
    client, engine, users, current, monkey = api
    group = '?factory_id=group'
    with Session(engine) as db:
        db.get(EmployeeProfile, 'stranger').confirmation_status = 'needs_review'
        db.commit()
    users['stranger'] = replace(users['stranger'], profile=replace(users['stranger'].profile, confirmation_status='needs_review'))
    assert 'stranger' not in {p['id'] for p in client.get(BASE + '/recipients' + group).json()['users']}
    doc = client.post(BASE + group, json={'format': 'xlsx'}).json()['id']
    assert put_grants(client, f'{BASE}/{doc}/grants{group}', json={'grants': [{'user_id': 'stranger', 'access': 'view'}]}).status_code == 422
    current['user'] = 'stranger'
    assert client.get(BASE + '/recipients' + group).status_code == 403
    own = client.post(BASE + group, json={'format': 'xlsx'})
    assert own.status_code == 201  # Personal files remain available.
    assert put_grants(client, f"{BASE}/{own.json()['id']}/grants{group}", json={'grants': []}).status_code == 403


@pytest.mark.parametrize('factory', ['group', 'huaxing'])
def test_stale_sharing_window_cannot_restore_revoked_grants(api, factory):
    client, engine, users, current, monkey = api
    suffix = '?factory_id=' + factory
    doc = client.post(BASE + suffix, json={'format': 'xlsx'}).json()['id']
    path = f'{BASE}/{doc}'
    grants = [{'user_id': 'viewer', 'access': 'view', 'epoch': 1}]
    original = put_grants(client, path + '/grants' + suffix, json={'grants': grants}).json()
    assert put_grants(client, path + '/grants' + suffix, json={'grants': []}).status_code == 200
    stale = client.put(path + '/grants' + suffix, json={'grants': grants, 'version': original['share_version']})
    assert stale.status_code == 409
    assert client.get(path + suffix).json()['grants'] == []


@pytest.mark.parametrize('factory', ['group', 'huaxing'])
def test_shared_grant_cannot_silently_rebind_to_rehired_employee(api, factory):
    client, engine, users, current, monkey = api
    suffix = '?factory_id=' + factory
    doc = client.post(BASE + suffix, json={'format': 'xlsx'}).json()['id']
    path = f'{BASE}/{doc}'
    grants = [{'user_id': 'viewer', 'access': 'view', 'epoch': 1}]
    assert put_grants(client, path + '/grants' + suffix, json={'grants': grants}).status_code == 200
    users['viewer'] = context('viewer', epoch=2)
    with Session(engine) as db:
        db.get(EmployeeProfile, 'viewer').employment_epoch = 2
        db.commit()
    current['user'] = 'viewer'
    assert client.get(path + suffix).status_code == 404
    current['user'] = 'owner'
    unchanged = client.get(path + suffix).json()['grants']
    result = put_grants(client, path + '/grants' + suffix, json={'grants': [*unchanged, {'user_id': 'editor', 'access': 'edit'}]})
    assert result.status_code == 409
    assert client.get(path + suffix).json()['grants'] == unchanged
    # The owner must explicitly reselect the current identity.
    assert put_grants(client, path + '/grants' + suffix, json={'grants': [{'user_id': 'viewer', 'access': 'view', 'epoch': 2}]}).status_code == 200
    current['user'] = 'viewer'
    assert client.get(path + suffix).json()['access'] == 'view'
