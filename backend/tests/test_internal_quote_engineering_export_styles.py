import json
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.services.internal_quote_excel import (
    ENGINEERING_WORKBOOK_TEMPLATE_PATH,
    build_internal_quote_engineering_workbook,
)


@pytest.mark.parametrize('count', [0, 1, 22, 23, 24, 25, 40, 120])
def test_purchase_detail_format_is_uniform_before_and_after_template_overflow(count):
    materials = [
        {'item': f'测试物料 {index}', 'category': 'hardware', 'specification': f'TEST-{index}',
         'quantity': index / 10, 'supplier': '测试供应商', 'color': '原色', 'material': '测试材质'}
        for index in range(1, count + 1)
    ]
    payload = json.dumps({'materials': materials}, ensure_ascii=False)
    section = SimpleNamespace(department='engineering', payload_json=payload)
    quote = SimpleNamespace(factory_id='huadeng', workshop_name='格式验收', customer='测试客户',
                            quote_no='BOM-STYLE-TEST', product_name='格式验收', version_label='V1', header_revision=1)
    template = load_workbook(ENGINEERING_WORKBOOK_TEMPLATE_PATH)
    output = load_workbook(BytesIO(build_internal_quote_engineering_workbook(quote, [section])))
    try:
        source = template['外购清单']
        sheet = output['外购清单']
        capacity = max(22, count)
        for row in range(6, capacity + 6):
            assert sheet.row_dimensions[row].height == source.row_dimensions[6].height
            for column in range(1, 11):
                assert sheet.cell(row, column)._style == source.cell(6, column)._style
        for index, material in enumerate(materials, start=1):
            assert sheet.cell(index + 5, 1).value == index
            assert sheet.cell(index + 5, 2).value == material['item']
            assert sheet.cell(index + 5, 3).value == material['specification']
            assert sheet.cell(index + 5, 6).value == material['quantity']
            assert sheet.cell(index + 5, 7).value == material['supplier']
        for row in range(count + 6, capacity + 6):
            assert all(sheet.cell(row, column).value is None for column in range(1, 11))
        for row in range(1, 6):
            for column in range(1, 11):
                assert sheet.cell(row, column)._style == source.cell(row, column)._style
            assert sheet.row_dimensions[row].height == source.row_dimensions[row].height
        for column in range(1, 11):
            letter = get_column_letter(column)
            assert sheet.column_dimensions[letter].width == source.column_dimensions[letter].width
        assert {str(value) for value in sheet.merged_cells.ranges} == {
            str(value) for value in source.merged_cells.ranges
        }
        assert str(sheet.print_area) == f"'外购清单'!$A$1:$J${capacity + 7}"
        assert output.sheetnames == ['排摸表', '外购清单']
        assert section.payload_json == payload
    finally:
        output.close()
        template.close()
