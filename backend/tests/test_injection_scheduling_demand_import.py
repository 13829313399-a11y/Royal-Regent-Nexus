from __future__ import annotations

import hashlib
import importlib
import json
import os
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select
from test_injection_scheduling_phase3_api import (
    ADMIN_TEST_PASSWORD,
    make_migrated_client,
)
from test_injection_scheduling_phase4_import import login


def _demand_workbook(*, row_count: int = 1, unknown_quantity_header: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "客户下单"
    sheet["A1"] = "啤机部生产啤货表"
    sheet["A2"] = "公司名称"
    sheet["B2"] = "测试客户"
    sheet["D2"] = "单号"
    sheet["E2"] = "SO-0001"
    sheet["G2"] = "交货日期:2026-08-20"
    sheet["I2"] = "交货地点"
    sheet["J2"] = "测试交货地点"
    headers = [
        "款号",
        "模具编号",
        "工模名称",
        "订单数量新口径" if unknown_quantity_header else "数量",
        "总套数",
        "啤数",
        "颜色",
        "色粉号",
        "用料名称",
        "模具日产量",
        "整啤毛重",
        "整啤净重",
        "总毛重",
        "总净重",
        "水口比例",
        "啤机复期",
        "备注",
    ]
    for column, value in enumerate(headers, start=1):
        sheet.cell(3, column, value)
    values = (
        "STYLE-01",
        "M-001",
        "测试工模",
        1000,
        100,
        500,
        "红色",
        "R-01",
        "ABS",
        2000,
        20,
        18,
        20,
        18,
        "10%",
        "待确认",
        "首批",
    )
    for row_index in range(row_count):
        row_values = list(values)
        if row_index:
            row_values[2] = f"测试工模-{row_index + 1}"
        for column, value in enumerate(row_values, start=1):
            sheet.cell(4 + row_index, column, value)
    sheet.cell(4 + row_count, 1, "备注")
    sheet.cell(6 + row_count, 1, "下单日期:2026-08-19")
    sheet.cell(6 + row_count, 10, "下单人:测试下单人")
    payload = BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def _master_data_workbook() -> bytes:
    workbook = Workbook()
    mold_sheet = workbook.active
    mold_sheet.title = "机安"
    mold_headers = [
        "工模编号",
        "工模名称",
        "货号",
        "配件名称",
        "安数",
        "模具安数",
        "日产量",
        "整啤毛重",
        "整啤净重",
    ]
    for column, value in enumerate(mold_headers, start=1):
        mold_sheet.cell(1, column, value)
    mold_values = ["M-002", "第二套测试工模", "STYLE-02", "上壳", "120A", 2, 1600, 24, 21]
    for column, value in enumerate(mold_values, start=1):
        mold_sheet.cell(2, column, value)

    price_sheet = workbook.create_sheet("单价")
    price_headers = ["模具编号", "货号", "产品名称", "单价/啤", "计价单位", "币种"]
    for column, value in enumerate(price_headers, start=1):
        price_sheet.cell(1, column, value)
    price_values = ["M-002", "STYLE-02", "上壳", "1.25", "PER_SHOT", "HKD"]
    for column, value in enumerate(price_values, start=1):
        price_sheet.cell(2, column, value)
    payload = BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def _seed_shared_resolution_data() -> None:
    db_module = importlib.import_module("app.db")
    shared = importlib.import_module("app.models.injection_scheduling_shared")
    normalize_identifier = importlib.import_module(
        "app.services.injection_scheduling_demand_import"
    ).normalize_identifier
    timestamp = "2026-08-09T12:00:00+08:00"
    namespace = "demand-order:company:huaxing"
    with db_module.SessionLocal() as db:
        customer = shared.InjectionSchedulingCustomerIdentity(
            id="customer:test",
            company_scope_id="company:royal-regent",
            customer_code="TEST",
            display_name="测试客户",
            status="ACTIVE",
            revision=1,
            created_at=timestamp,
        )
        definition = shared.InjectionSchedulingMoldDefinition(
            id="mold-definition:test-001",
            company_scope_id="company:royal-regent",
            canonical_mold_no="M-001",
            display_mold_no="M-001",
            standard_name="测试工模",
            status="ACTIVE",
            data_quality="APPROVED",
            revision=1,
            valid_from="2026-01-01",
            created_factory_id="huaxing",
            approved_by="system-test",
            created_at=timestamp,
            approved_at=timestamp,
        )
        output = shared.InjectionSchedulingMoldOutputSpec(
            id="mold-output:test-001-style-01",
            mold_definition_id=definition.id,
            customer_identity_id=None,
            item_no="STYLE-01",
            product_name="测试工模",
            status="ACTIVE",
            revision=1,
            valid_from="2026-01-01",
            created_at=timestamp,
        )
        db.add_all([customer, definition])
        db.flush()
        db.add_all(
            [
                shared.InjectionSchedulingCustomerAlias(
                    id="customer-alias:test",
                    company_scope_id="company:royal-regent",
                    customer_identity_id=customer.id,
                    source_namespace_id=namespace,
                    raw_alias="测试客户",
                    normalized_alias=normalize_identifier("测试客户"),
                    valid_from="2026-01-01",
                    status="ACTIVE",
                    revision=1,
                    created_at=timestamp,
                ),
                shared.InjectionSchedulingMoldAlias(
                    id="mold-alias:test-001",
                    company_scope_id="company:royal-regent",
                    source_namespace_id=namespace,
                    namespace_type="FACTORY_SOURCE",
                    raw_alias="M-001",
                    normalized_alias=normalize_identifier("M-001"),
                    mold_definition_id=definition.id,
                    status="ACTIVE",
                    revision=1,
                    valid_from="2026-01-01",
                    created_at=timestamp,
                ),
                output,
            ]
        )
        db.flush()
        db.add(
            shared.InjectionSchedulingCommercialRateRule(
                    id="rate:test-001",
                    owner_scope_type="COMPANY",
                    owner_scope_id="company:royal-regent",
                    applicable_factory_mode="ANY_IN_OWNER_COMPANY",
                    applicable_factory_id=None,
                    applicable_customer_mode="ANY",
                    applicable_customer_id=None,
                    mold_definition_id=None,
                    mold_output_spec_id=output.id,
                    contract_mode="ANY",
                    contract_id=None,
                    pricing_basis="PER_SHOT",
                    amount=1,
                    currency="HKD",
                    tax_mode="EXCLUDED",
                    priority=10,
                    status="ACTIVE",
                    revision=1,
                    valid_from="2026-01-01",
                    proposed_by="system-test-proposer",
                    approved_by="system-test-approver",
                    created_at=timestamp,
                    approved_at=timestamp,
            )
        )
        db.commit()


def test_demand_order_preview_and_confirm_only_create_draft_backlog(monkeypatch):
    with make_migrated_client(monkeypatch) as client:
        login(client, "admin", ADMIN_TEST_PASSWORD)
        _seed_shared_resolution_data()
        shared_search = client.get(
            "/api/injection-scheduling/shared-molds/definitions",
            params={"factory_id": "huakang-b", "q": "M-001"},
        )
        assert shared_search.status_code == 200, shared_search.text
        assert shared_search.json()[0]["canonical_mold_no"] == "M-001"
        assert (
            shared_search.json()[0]["factory_readiness"]["status"]
            == "NOT_FACTORY_READY"
        )
        preview_response = client.post(
            "/api/injection-scheduling/imports/preview",
            data={
                "factory_id": "huaxing",
                "expected_revision": "0",
                "document_kind": "DEMAND_ORDER",
            },
            files={
                "file": (
                    "demand-order.xlsx",
                    _demand_workbook(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "demand-preview-0001"},
        )
        assert preview_response.status_code == 201, preview_response.text
        preview = preview_response.json()
        assert preview["document_kind"] == "DEMAND_ORDER"
        assert preview["batch_state"] == "PREVIEW_READY"
        assert preview["summary"]["ready_count"] == 1
        canonical = preview["demand_rows"][0]["canonical"]
        assert canonical["delivery_due_date"] == "2026-08-20"
        assert canonical["warehouse_text"] == "测试下单人"
        assert canonical["order_date"] == "2026-08-19"
        assert canonical["product_group_no"] == "STYLE-01"
        assert canonical["source_mold_no"] == "M-001"
        assert canonical["product_name"] == "测试工模"
        assert canonical["order_quantity"] == 1000
        assert canonical["color_name"] == "红色"
        assert canonical["color_powder_code"] == "R-01"
        assert canonical["material_name"] == "ABS"
        assert canonical["source_daily_capacity"] == 2000
        assert canonical["whole_shot_net_weight_g"] == 18
        assert canonical["total_gross_weight"] == 20
        assert canonical["sprue_ratio"] == 0.1
        assert canonical["remark"] == "首批"
        assert (
            preview["demand_rows"][0]["resolved_values"]["factory_readiness_status"]
            == "NOT_FACTORY_READY"
        )

        confirm_response = client.post(
            f"/api/injection-scheduling/imports/{preview['id']}/confirm",
            json={
                "factory_id": "huaxing",
                "document_kind": "DEMAND_ORDER",
                "expected_revision": 1,
                "expected_plan_revision": 0,
                "expected_preview_generation": 1,
                "expected_resolution_digest": preview["resolution_digest"],
                "request_id": "demand-confirm-0001",
                "confirm_mode": "create_draft",
                "confirm_scope": "ALL_READY",
                "business_date": "2026-08-09",
            },
        )
        assert confirm_response.status_code == 200, confirm_response.text
        confirmed = confirm_response.json()
        assert confirmed["status"] == "CONFIRMED"
        assert confirmed["result"]["created_tasks"] == 0
        assert confirmed["result"]["backlog_only"] is True

        db_module = importlib.import_module("app.db")
        execution = importlib.import_module("app.models.injection_scheduling_execution")
        shared = importlib.import_module("app.models.injection_scheduling_shared")
        with db_module.SessionLocal() as db:
            assert (
                db.scalar(
                    select(func.count()).select_from(execution.InjectionSchedulingTask)
                )
                == 0
            )
            state = db.scalar(select(execution.InjectionSchedulingPlanOrderState))
            assert state is not None
            assert state.status == "BACKLOG"
            assert state.factory_readiness_status == "NOT_FACTORY_READY"
            assert state.order_revision_id
            assert db.get(
                shared.InjectionSchedulingDemandOrderVersion, state.order_revision_id
            )
            order = db.get(execution.InjectionSchedulingOrder, state.order_id)
            assert order is not None
            assert order.order_no == "SO-0001"
            assert order.item_no == "STYLE-01"
            assert order.product_name == "测试工模"
            assert order.order_quantity == Decimal(1000)
            assert order.delivery_due_date == "2026-08-20"
            assert order.warehouse_text == "测试下单人"
            assert order.remark == "首批"
            lineage = json.loads(order.lineage_json)
            assert lineage["source_mold_no"] == "M-001"
            assert lineage["source_daily_capacity"] == 2000
            assert lineage["color_name"] == "红色"
            assert lineage["color_powder_code"] == "R-01"
            assert lineage["material_name"] == "ABS"
            assert lineage["whole_shot_net_weight_g"] == 18
            assert lineage["total_gross_weight"] == 20
            assert lineage["sprue_ratio"] == 0.1
            assert lineage["order_date"] == "2026-08-19"
            default_shift_target = importlib.import_module(
                "app.services.injection_scheduling_scheduler.duration"
            ).default_shift_target
            assert default_shift_target(order) == Decimal(1000)

        mapping_response = client.post(
            "/api/injection-scheduling/imports/preview",
            data={
                "factory_id": "huaxing",
                "expected_revision": "0",
                "document_kind": "DEMAND_ORDER",
            },
            files={
                "file": (
                    "unknown-header.xlsx",
                    _demand_workbook(unknown_quantity_header=True),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "demand-preview-mapping-0001"},
        )
        assert mapping_response.status_code == 201, mapping_response.text
        mapping_preview = mapping_response.json()
        assert mapping_preview["batch_state"] == "MAPPING_REQUIRED"
        assert any(
            item["canonical_field"] == "order_quantity"
            and item["status"] == "MISSING"
            for item in mapping_preview["mapping"]
        )
        saved_mapping = client.patch(
            f"/api/injection-scheduling/imports/{mapping_preview['id']}/mapping-draft",
            json={
                "factory_id": "huaxing",
                "expected_revision": 1,
                "request_id": "demand-mapping-save-0001",
                "mappings": {"order_quantity": "订单数量新口径"},
            },
        )
        assert saved_mapping.status_code == 200, saved_mapping.text
        proposed_profile = client.post(
            f"/api/injection-scheduling/imports/{mapping_preview['id']}/profile-proposal",
            json={
                "factory_id": "huaxing",
                "expected_revision": saved_mapping.json()["revision"],
                "request_id": "demand-profile-proposal-0001",
                "name": "下单表数量字段别名修订",
                "reason": "客户新模板将数量表头改名",
            },
        )
        assert proposed_profile.status_code == 200, proposed_profile.text
        assert proposed_profile.json()["batch_state"] == "PROFILE_REVIEW_PENDING"
        assert proposed_profile.json()["mapping_draft"]["proposal_profile_id"]

        workbench_mapping_response = client.post(
            "/api/injection-scheduling/imports/preview",
            data={
                "factory_id": "huaxing",
                "expected_revision": "0",
                "document_kind": "DEMAND_ORDER",
            },
            files={
                "file": (
                    "unknown-header-workbench.xlsx",
                    _demand_workbook(unknown_quantity_header=True),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "workbench-demand-preview-0001"},
        )
        assert workbench_mapping_response.status_code == 201
        workbench_mapping_preview = workbench_mapping_response.json()
        assert workbench_mapping_preview["batch_state"] == "MAPPING_REQUIRED"

        applied_mapping = client.post(
            "/api/injection-scheduling/workbench/imports/"
            f"{workbench_mapping_preview['id']}/apply-mapping",
            json={
                "factory_id": "huaxing",
                "expected_revision": workbench_mapping_preview["revision"],
                "request_id": "workbench-demand-mapping-0001",
                "mappings": {"order_quantity": "订单数量新口径"},
            },
        )
        assert applied_mapping.status_code == 200, applied_mapping.text
        corrected_preview = applied_mapping.json()
        assert corrected_preview["batch_state"] == "PREVIEW_READY"
        assert corrected_preview["preview_generation"] == 2
        assert corrected_preview["demand_rows"][0]["canonical"]["order_quantity"] == 1000

        mapping_replay = client.post(
            "/api/injection-scheduling/workbench/imports/"
            f"{workbench_mapping_preview['id']}/apply-mapping",
            json={
                "factory_id": "huaxing",
                "expected_revision": workbench_mapping_preview["revision"],
                "request_id": "workbench-demand-mapping-0001",
                "mappings": {"order_quantity": "订单数量新口径"},
            },
        )
        assert mapping_replay.status_code == 200, mapping_replay.text
        assert mapping_replay.json()["revision"] == corrected_preview["revision"]

        workbench_confirm = client.post(
            f"/api/injection-scheduling/imports/{corrected_preview['id']}/confirm",
            json={
                "factory_id": "huaxing",
                "document_kind": "DEMAND_ORDER",
                "expected_revision": corrected_preview["revision"],
                "expected_plan_revision": corrected_preview["plan_context"][
                    "target_draft_plan_revision"
                ],
                "expected_preview_generation": corrected_preview[
                    "preview_generation"
                ],
                "expected_resolution_digest": corrected_preview[
                    "resolution_digest"
                ],
                "request_id": "workbench-demand-confirm-0001",
                "confirm_mode": "merge_draft",
                "confirm_scope": "SELECTED",
                "selected_row_ids": [corrected_preview["demand_rows"][0]["row_id"]],
                "business_date": "2026-08-09",
            },
        )
        assert workbench_confirm.status_code == 200, workbench_confirm.text
        assert workbench_confirm.json()["status"] == "CONFIRMED"

        import_models = importlib.import_module(
            "app.models.injection_scheduling_import"
        )
        profile_registry = importlib.import_module(
            "app.services.injection_scheduling_profile_registry"
        )
        with db_module.SessionLocal() as db:
            profile_registry.seed_builtin_import_profiles(db)
            workbench_profile = db.scalar(
                select(import_models.InjectionSchedulingImportProfile)
                .where(
                    import_models.InjectionSchedulingImportProfile.profile_family
                    == "demand_order_shared",
                    import_models.InjectionSchedulingImportProfile.status == "ACTIVE",
                    import_models.InjectionSchedulingImportProfile.profile_code.like(
                        "%demand_order_shared-wb-r%"
                    ),
                )
                .order_by(
                    import_models.InjectionSchedulingImportProfile.revision.desc()
                )
            )
            assert workbench_profile is not None
            assert "订单数量新口径" in workbench_profile.config_json
            huaxing_demand_profiles = profile_registry.active_profiles_for_factory(
                db, "huaxing", "DEMAND_ORDER"
            )
            other_factory_profiles = profile_registry.active_profiles_for_factory(
                db, "huakang-b", "DEMAND_ORDER"
            )
            assert [item.profile_id for item in huaxing_demand_profiles] == [
                workbench_profile.id
            ]
            assert [item.profile_id for item in other_factory_profiles] == [
                "isprofile-demand-order-shared-v1"
            ]

        partial_response = client.post(
            "/api/injection-scheduling/imports/preview",
            data={
                "factory_id": "huaxing",
                "expected_revision": "0",
                "document_kind": "DEMAND_ORDER",
            },
            files={
                "file": (
                    "two-demand-rows.xlsx",
                    _demand_workbook(row_count=2),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "demand-preview-partial-0001"},
        )
        assert partial_response.status_code == 201, partial_response.text
        partial = partial_response.json()
        assert partial["summary"]["ready_count"] == 2
        first_row_id, second_row_id = [item["row_id"] for item in partial["demand_rows"]]
        partial_confirm = client.post(
            f"/api/injection-scheduling/imports/{partial['id']}/confirm",
            json={
                "factory_id": "huaxing",
                "document_kind": "DEMAND_ORDER",
                "expected_revision": 1,
                "expected_plan_revision": partial["plan_context"][
                    "target_draft_plan_revision"
                ],
                "expected_preview_generation": 1,
                "expected_resolution_digest": partial["resolution_digest"],
                "request_id": "demand-confirm-partial-0001",
                "confirm_mode": "merge_draft",
                "confirm_scope": "SELECTED",
                "selected_row_ids": [first_row_id],
                "business_date": "2026-08-09",
            },
        )
        assert partial_confirm.status_code == 200, partial_confirm.text
        partially_confirmed = partial_confirm.json()
        assert partially_confirmed["status"] == "PARTIALLY_CONFIRMED"
        assert [item["confirmation_state"] for item in partially_confirmed["demand_rows"]] == [
            "CONFIRMED",
            "PENDING",
        ]
        final_confirm = client.post(
            f"/api/injection-scheduling/imports/{partial['id']}/confirm",
            json={
                "factory_id": "huaxing",
                "document_kind": "DEMAND_ORDER",
                "expected_revision": partially_confirmed["revision"],
                "expected_plan_revision": partially_confirmed[
                    "confirmed_plan_revision"
                ],
                "expected_preview_generation": 1,
                "expected_resolution_digest": partial["resolution_digest"],
                "request_id": "demand-confirm-partial-0002",
                "confirm_mode": "merge_draft",
                "confirm_scope": "SELECTED",
                "selected_row_ids": [second_row_id],
                "target_draft_plan_id": partially_confirmed["confirmed_plan_id"],
                "business_date": "2026-08-09",
            },
        )
        assert final_confirm.status_code == 200, final_confirm.text
        assert final_confirm.json()["status"] == "CONFIRMED"
        assert final_confirm.json()["result"]["created_tasks"] == 0

        stale_response = client.post(
            "/api/injection-scheduling/imports/preview",
            data={
                "factory_id": "huaxing",
                "expected_revision": "0",
                "document_kind": "DEMAND_ORDER",
            },
            files={
                "file": (
                    "stale-demand.xlsx",
                    _demand_workbook(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "demand-preview-stale-0001"},
        )
        assert stale_response.status_code == 201, stale_response.text
        stale = stale_response.json()
        with db_module.SessionLocal() as db:
            alias = db.get(shared.InjectionSchedulingMoldAlias, "mold-alias:test-001")
            alias.revision += 1
            db.commit()
        stale_confirm = client.post(
            f"/api/injection-scheduling/imports/{stale['id']}/confirm",
            json={
                "factory_id": "huaxing",
                "document_kind": "DEMAND_ORDER",
                "expected_revision": 1,
                "expected_plan_revision": stale["plan_context"][
                    "target_draft_plan_revision"
                ],
                "expected_preview_generation": 1,
                "expected_resolution_digest": stale["resolution_digest"],
                "request_id": "demand-confirm-stale-0001",
                "confirm_mode": "merge_draft",
                "confirm_scope": "ALL_READY",
                "business_date": "2026-08-09",
            },
        )
        assert stale_confirm.status_code == 409, stale_confirm.text
        assert stale_confirm.json()["detail"]["code"] == "MASTER_DATA_STALE"


def test_customer_and_price_are_not_scheduling_import_prerequisites(monkeypatch):
    with make_migrated_client(monkeypatch) as client:
        login(client, "admin", ADMIN_TEST_PASSWORD)
        _seed_shared_resolution_data()

        db_module = importlib.import_module("app.db")
        shared = importlib.import_module("app.models.injection_scheduling_shared")
        with db_module.SessionLocal() as db:
            db.query(shared.InjectionSchedulingCustomerAlias).delete()
            db.query(shared.InjectionSchedulingCommercialRateRule).delete()
            db.commit()

        preview_response = client.post(
            "/api/injection-scheduling/imports/preview",
            data={
                "factory_id": "huaxing",
                "expected_revision": "0",
                "document_kind": "DEMAND_ORDER",
            },
            files={
                "file": (
                    "demand-without-customer-master.xlsx",
                    _demand_workbook(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "demand-preview-no-customer-0001"},
        )
        assert preview_response.status_code == 201, preview_response.text
        preview = preview_response.json()
        assert preview["summary"]["ready_count"] == 1
        row = preview["demand_rows"][0]
        assert row["resolution_status"] == "READY"
        assert row["canonical"]["customer_name"] == "测试客户"
        assert row["resolved_values"]["customer_identity_id"] is None
        assert row["resolved_values"]["commercial_rate_rule_id"] is None
        assert row["resolved_values"]["frozen_amount"] is None

        confirm_response = client.post(
            f"/api/injection-scheduling/imports/{preview['id']}/confirm",
            json={
                "factory_id": "huaxing",
                "document_kind": "DEMAND_ORDER",
                "expected_revision": 1,
                "expected_plan_revision": 0,
                "expected_preview_generation": 1,
                "expected_resolution_digest": preview["resolution_digest"],
                "request_id": "demand-confirm-no-customer-0001",
                "confirm_mode": "create_draft",
                "confirm_scope": "ALL_READY",
                "business_date": "2026-08-09",
            },
        )
        assert confirm_response.status_code == 200, confirm_response.text
        with db_module.SessionLocal() as db:
            version = db.scalar(
                select(shared.InjectionSchedulingDemandOrderVersion)
            )
            assert version is not None
            assert version.customer_identity_id is None
            assert version.commercial_rate_rule_id is None
            assert "测试客户" in version.canonical_json


def test_missing_mold_master_enters_backlog_for_later_enrichment(monkeypatch):
    with make_migrated_client(monkeypatch) as client:
        login(client, "admin", ADMIN_TEST_PASSWORD)
        preview_response = client.post(
            "/api/injection-scheduling/imports/preview",
            data={
                "factory_id": "huaxing",
                "expected_revision": "0",
                "document_kind": "DEMAND_ORDER",
            },
            files={
                "file": (
                    "demand-before-mold-master.xlsx",
                    _demand_workbook(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "demand-preview-before-mold-0001"},
        )
        assert preview_response.status_code == 201, preview_response.text
        preview = preview_response.json()
        assert preview["summary"]["ready_count"] == 1
        row = preview["demand_rows"][0]
        assert row["resolution_status"] == "READY"
        assert row["resolved_values"]["mold_definition_id"] is None
        assert row["resolved_values"]["mold_enrichment_status"] == "PENDING"
        assert row["resolved_values"]["factory_readiness_status"] == "NOT_FACTORY_READY"
        assert row["resolution_reasons"] == ["shared_mold_pending_enrichment"]

        confirm_response = client.post(
            f"/api/injection-scheduling/imports/{preview['id']}/confirm",
            json={
                "factory_id": "huaxing",
                "document_kind": "DEMAND_ORDER",
                "expected_revision": 1,
                "expected_plan_revision": 0,
                "expected_preview_generation": 1,
                "expected_resolution_digest": preview["resolution_digest"],
                "request_id": "demand-confirm-before-mold-0001",
                "confirm_mode": "create_draft",
                "confirm_scope": "ALL_READY",
                "business_date": "2026-08-09",
            },
        )
        assert confirm_response.status_code == 200, confirm_response.text

        db_module = importlib.import_module("app.db")
        master = importlib.import_module("app.models.injection_scheduling")
        execution = importlib.import_module("app.models.injection_scheduling_execution")
        with db_module.SessionLocal() as db:
            order = db.scalar(select(execution.InjectionSchedulingOrder))
            assert order is not None
            assert order.status == "BACKLOG"
            assert order.mold_id is None
            assert '"source_mold_no":"M-001"' in order.lineage_json
            state = db.scalar(select(execution.InjectionSchedulingPlanOrderState))
            assert state is not None
            assert state.factory_readiness_status == "NOT_FACTORY_READY"
            db.add(
                master.InjectionSchedulingMold(
                    id="mold:late-master-001",
                    factory_id="huaxing",
                    mold_no="M-001",
                    name="后补测试工模",
                    copy_count=1,
                    data_quality_status="complete",
                    status="available",
                    revision=1,
                    normalization_status="COMPLETE",
                    created_by="system-test",
                    created_by_name="测试",
                    updated_by="system-test",
                    updated_by_name="测试",
                    created_at="2026-08-09T12:00:00+08:00",
                    updated_at="2026-08-09T12:00:00+08:00",
                )
            )
            db.commit()

        enrichment_response = client.post(
            "/api/injection-scheduling/demand-orders/enrich-molds",
            json={
                "factory_id": "huaxing",
                "request_id": "demand-mold-enrichment-0001",
                "reason": "测试导入后从模具库补齐",
            },
        )
        assert enrichment_response.status_code == 200, enrichment_response.text
        assert enrichment_response.json()["matched"] == 1
        assert enrichment_response.json()["pending"] == 0
        with db_module.SessionLocal() as db:
            order = db.scalar(select(execution.InjectionSchedulingOrder))
            assert order is not None
            assert order.mold_id == "mold:late-master-001"
            assert '"mold_enrichment_status":"MATCHED"' in order.lineage_json
            state = db.scalar(select(execution.InjectionSchedulingPlanOrderState))
            assert state is not None
            assert state.factory_readiness_status == "NOT_FACTORY_READY"


def test_shared_master_enrichment_does_not_require_legacy_or_physical_mold(
    monkeypatch,
):
    with make_migrated_client(monkeypatch) as client:
        login(client, "admin", ADMIN_TEST_PASSWORD)
        preview_response = client.post(
            "/api/injection-scheduling/imports/preview",
            data={
                "factory_id": "huaxing",
                "expected_revision": "0",
                "document_kind": "DEMAND_ORDER",
            },
            files={
                "file": (
                    "demand-before-shared-master.xlsx",
                    _demand_workbook(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "demand-preview-shared-master-0001"},
        )
        assert preview_response.status_code == 201, preview_response.text
        preview = preview_response.json()
        confirm_response = client.post(
            f"/api/injection-scheduling/imports/{preview['id']}/confirm",
            json={
                "factory_id": "huaxing",
                "document_kind": "DEMAND_ORDER",
                "expected_revision": 1,
                "expected_plan_revision": 0,
                "expected_preview_generation": 1,
                "expected_resolution_digest": preview["resolution_digest"],
                "request_id": "demand-confirm-shared-master-0001",
                "confirm_mode": "create_draft",
                "confirm_scope": "ALL_READY",
                "business_date": "2026-08-09",
            },
        )
        assert confirm_response.status_code == 200, confirm_response.text

        db_module = importlib.import_module("app.db")
        execution = importlib.import_module("app.models.injection_scheduling_execution")
        shared = importlib.import_module("app.models.injection_scheduling_shared")
        timestamp = "2026-08-09T12:00:00+08:00"
        with db_module.SessionLocal() as db:
            definition = shared.InjectionSchedulingMoldDefinition(
                id="mold-definition:late-shared-001",
                company_scope_id="company:royal-regent",
                canonical_mold_no="M-001",
                display_mold_no="M-001",
                standard_name="共享测试工模",
                recommended_machine_class_raw="7A",
                mold_a_class=7,
                default_arm_type="SINGLE",
                default_fixture_type="SUCTION_CUP",
                status="ACTIVE",
                data_quality="APPROVED",
                revision=1,
                valid_from="2026-01-01",
                created_factory_id="huaxing",
                approved_by="system-test",
                created_at=timestamp,
                approved_at=timestamp,
            )
            output = shared.InjectionSchedulingMoldOutputSpec(
                id="mold-output:late-shared-001",
                mold_definition_id=definition.id,
                customer_identity_id=None,
                item_no="STYLE-01",
                product_name="共享测试工模",
                cavity_count=2,
                whole_shot_net_weight_g=18,
                nominal_daily_capacity=3600,
                status="ACTIVE",
                revision=1,
                valid_from="2026-01-01",
                created_at=timestamp,
            )
            capability = shared.InjectionSchedulingFactoryMoldCapability(
                id="mold-capability:late-shared-001",
                factory_id="huaxing",
                mold_definition_id=definition.id,
                mold_output_spec_id=output.id,
                physical_asset_id=None,
                machine_class=7,
                applicability_key="DEFAULT",
                required_arm_type="SINGLE",
                required_fixture_type="SUCTION_CUP",
                nominal_daily_capacity=3600,
                priority=100,
                status="ACTIVE",
                revision=1,
                valid_from="2026-01-01",
                approved_by="system-test",
                approved_at=timestamp,
                created_at=timestamp,
            )
            rate = shared.InjectionSchedulingCommercialRateRule(
                id="rate:late-shared-001",
                owner_scope_type="FACTORY",
                owner_scope_id="huaxing",
                applicable_factory_mode="SPECIFIC",
                applicable_factory_id="huaxing",
                applicable_customer_mode="ANY",
                applicable_customer_id=None,
                mold_definition_id=definition.id,
                mold_output_spec_id=output.id,
                contract_mode="ANY",
                contract_id=None,
                pricing_basis="PER_SHOT",
                amount=Decimal("1.25"),
                currency="CNY",
                tax_mode="AS_LISTED",
                priority=100,
                status="ACTIVE",
                revision=1,
                valid_from="2026-08-10T01:00:00+08:00",
                proposed_by="system-test-proposer",
                approved_by="system-test-approver",
                created_at=timestamp,
                approved_at=timestamp,
            )
            db.add(definition)
            db.flush()
            db.add(output)
            db.flush()
            db.add_all([capability, rate])
            db.commit()

        enrichment_response = client.post(
            "/api/injection-scheduling/demand-orders/enrich-molds",
            json={
                "factory_id": "huaxing",
                "request_id": "demand-shared-enrichment-0001",
                "reason": "测试直接从共享主数据补齐",
            },
        )
        assert enrichment_response.status_code == 200, enrichment_response.text
        result = enrichment_response.json()
        assert result["matched"] == 1
        assert result["pending"] == 0
        assert result["ambiguous"] == 0
        assert result["execution_ready"] == 0

        with db_module.SessionLocal() as db:
            order = db.scalar(select(execution.InjectionSchedulingOrder))
            assert order is not None
            assert order.mold_id is None
            lineage = json.loads(order.lineage_json)
            assert lineage["mold_enrichment_status"] == "MATCHED"
            assert lineage["mold_definition_id"] == definition.id
            assert lineage["mold_output_spec_id"] == output.id
            assert lineage["factory_capability_id"] == capability.id
            assert lineage["machine_a_class"] == 7
            assert lineage["required_arm_type"] == "SINGLE"
            assert lineage["required_fixture_type"] == "SUCTION_CUP"
            assert lineage["mold_daily_capacity"] == 3600
            assert lineage["mold_whole_shot_net_weight_g"] == 18
            assert lineage["unit_price"] == 1.25
            assert lineage["currency"] == "CNY"
            assert lineage["pricing_basis"] == "PER_SHOT"
            assert lineage["factory_readiness_status"] == "NOT_FACTORY_READY"
            state = db.scalar(select(execution.InjectionSchedulingPlanOrderState))
            assert state is not None
            assert state.factory_readiness_status == "NOT_FACTORY_READY"


def test_master_data_preview_only_creates_governed_proposals(monkeypatch):
    with make_migrated_client(monkeypatch) as client:
        login(client, "admin", ADMIN_TEST_PASSWORD)
        preview_response = client.post(
            "/api/injection-scheduling/imports/preview",
            data={
                "factory_id": "huaxing",
                "expected_revision": "0",
                "document_kind": "MASTER_DATA",
            },
            files={
                "file": (
                    "master-data.xlsx",
                    _master_data_workbook(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "master-preview-0001"},
        )
        assert preview_response.status_code == 201, preview_response.text
        preview = preview_response.json()
        assert preview["document_kind"] == "MASTER_DATA"
        assert preview["batch_state"] == "PREVIEW_READY"
        assert preview["summary"]["proposable_count"] == 2
        assert {item["entity_type"] for item in preview["master_data_rows"]} == {
            "MOLD_DEFINITION_BUNDLE",
            "COMMERCIAL_RATE_RULE",
        }

        confirm_response = client.post(
            f"/api/injection-scheduling/imports/{preview['id']}/confirm",
            json={
                "factory_id": "huaxing",
                "document_kind": "MASTER_DATA",
                "expected_revision": preview["revision"],
                "expected_plan_revision": 0,
                "expected_preview_generation": preview["preview_generation"],
                "expected_resolution_digest": preview["resolution_digest"],
                "request_id": "master-confirm-0001",
                "confirm_mode": "propose_master_data",
                "confirm_scope": "ALL_READY",
                "business_date": "2026-08-09",
            },
        )
        assert confirm_response.status_code == 200, confirm_response.text
        confirmed = confirm_response.json()
        assert confirmed["status"] == "CONFIRMED"
        assert confirmed["result"]["created_proposals"] == 2
        assert confirmed["result"]["activation_performed"] is False

        db_module = importlib.import_module("app.db")
        shared = importlib.import_module("app.models.injection_scheduling_shared")
        with db_module.SessionLocal() as db:
            proposals = list(
                db.scalars(select(shared.InjectionSchedulingMasterDataProposal)).all()
            )
            assert len(proposals) == 2
            assert {item.status for item in proposals} == {"PROPOSED"}
            assert db.scalar(
                select(func.count()).select_from(shared.InjectionSchedulingFieldEvidence)
            ) >= 10
            assert (
                db.scalar(
                    select(func.count()).select_from(
                        shared.InjectionSchedulingCommercialRateRule
                    )
                )
                == 0
            )


            mold_proposal = next(
                item
                for item in proposals
                if item.entity_type == "MOLD_DEFINITION_BUNDLE"
            )
            mold_proposal.status = "APPROVED"
            mold_proposal.revision = 2
            mold_proposal.approved_by = "independent-master-reviewer"
            mold_proposal.approved_by_name = "独立主数据审核员"
            mold_proposal.reviewed_at = "2026-08-09T12:30:00+08:00"
            db.commit()
            proposal_id = mold_proposal.id

        activate_response = client.post(
            f"/api/injection-scheduling/master-data-proposals/{proposal_id}/activate",
            json={
                "factory_id": "huaxing",
                "expected_revision": 2,
                "request_id": "master-activate-0001",
                "reason": "独立审核通过后激活公司模具定义",
            },
        )
        assert activate_response.status_code == 200, activate_response.text
        activated = activate_response.json()
        assert activated["status"] == "ACTIVE"
        assert activated["activated_entity"]["mold_definition_id"]

        definitions_response = client.get(
            "/api/injection-scheduling/shared-molds/definitions",
            params={"factory_id": "huakang-b", "q": "M-002"},
        )
        assert definitions_response.status_code == 200, definitions_response.text
        assert [item["canonical_mold_no"] for item in definitions_response.json()] == [
            "M-002"
        ]
        with db_module.SessionLocal() as db:
            assert (
                db.scalar(
                    select(func.count()).select_from(
                        shared.InjectionSchedulingCommercialRateRule
                    )
                )
                == 0
            )


@pytest.mark.skipif(
    not os.getenv("INJECTION_SCHEDULING_DEMAND_WORKBOOK"),
    reason="set INJECTION_SCHEDULING_DEMAND_WORKBOOK for local demand-template regression",
)
def test_real_demand_workbook_read_only_preview(monkeypatch):
    source_path = Path(os.environ["INJECTION_SCHEDULING_DEMAND_WORKBOOK"])
    source = source_path.read_bytes()
    before_hash = hashlib.sha256(source).hexdigest()
    before_stat = source_path.stat()
    with make_migrated_client(monkeypatch) as client:
        login(client, "admin", ADMIN_TEST_PASSWORD)
        response = client.post(
            "/api/injection-scheduling/imports/preview",
            data={
                "factory_id": "huaxing",
                "expected_revision": "0",
                "document_kind": "DEMAND_ORDER",
            },
            files={
                "file": (
                    source_path.name,
                    source,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers={"x-request-id": "real-demand-read-only-preview"},
        )
        assert response.status_code == 201, response.text
        payload = response.json()
        assert payload["document_kind"] == "DEMAND_ORDER"
        assert payload["source_file_hash"] == before_hash
        assert payload["demand_rows"]
        assert all(
            "machine_code" not in row["canonical"]
            for row in payload["demand_rows"]
        )
    after_stat = source_path.stat()
    assert hashlib.sha256(source_path.read_bytes()).hexdigest() == before_hash
    assert after_stat.st_size == before_stat.st_size
    assert after_stat.st_mtime_ns == before_stat.st_mtime_ns
