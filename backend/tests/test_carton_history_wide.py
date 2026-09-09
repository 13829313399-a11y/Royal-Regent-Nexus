from decimal import Decimal
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from app.services.carton_procurement_history_import import _parse_rows, _group_rows


HEADERS = ["客户名称", "下单日期", "合同号", "货号", "产品订单数量", "内箱需求数量", "外箱需求数量",
    "计划交期", "客户交期", "备注", "产品名称", "内箱纸质", "内箱规格", "内箱每箱个数",
    "外箱纸质", "外箱规格", "外箱每箱个数", "历史订单号", "装箱数"]


def workbook(rows, headers=HEADERS, extras=None):
    wb=Workbook(); ws=wb.active; ws.title="历史订单导入"
    ws.append(["历史需求：不会建立库存"]); ws.append(headers)
    for row in rows:
        ws.append([row.get(key) for key in headers] if isinstance(row,dict) else row)
    if extras is not None:
        ws=wb.create_sheet("附加纸品明细")
        ws.append(["历史订单号","纸品类型","纸品需求数量","纸质","规格","纸品单位"])
        for row in extras: ws.append(row)
    example=wb.create_sheet("填写示例");example.append(headers);example.append(["不导入"])
    stream=BytesIO();wb.save(stream);return stream.getvalue()


def row(**kwargs):
    return {"客户名称":"迪奇","下单日期":"2026-09-02","合同号":"4500219760-10","货号":"92120TQ1-S001-NB",
        "外箱需求数量":100,"计划交期":"2026-09-10","外箱纸质":"K3K","外箱规格":"44.45*22.54*39.05",**kwargs}


def test_explicit_outer_demand_and_zero_inner_preserve_unknown_product_quantity():
    rows,warnings=_parse_rows("history.xlsx",workbook([row(装箱数="0/12")]))
    assert len(rows)==1
    assert rows[0]["product_order_quantity"] is None
    assert rows[0]["quantity_basis"]=="EXPLICIT"
    assert rows[0]["line"].usage_quantity==12
    assert rows[0]["line"].required_quantity==100
    assert rows[0]["line"].packaging_type=="外箱"


MINIMAL_HEADERS = ["客户名称", "合同号", "货号", "下单日期", "计划交期", "外箱需求数量", "内箱需求数量", "卡纸需求数量"]


def test_minimal_eight_columns_preserve_three_independent_demands_without_material_or_packing():
    rows,_=_parse_rows("history.xlsx",workbook([row(内箱需求数量=200,卡纸需求数量=300)],MINIMAL_HEADERS))
    assert len(_group_rows(rows))==1
    assert {r["line"].packaging_type:(r["line"].required_quantity,r["line"].unit) for r in rows}=={
        "外箱":(100,"个"),"内箱":(200,"个"),"卡纸":(300,"张")}
    assert all(r["product_order_quantity"] is None and r["line"].usage_quantity is None for r in rows)
    assert all(r["line"].paper_quality==r["line"].specification=="" for r in rows)
    assert any("按“卡纸”" in warning for r in rows for warning in r["row_warnings"])


@pytest.mark.parametrize("header",["卡纸需求数量","卡纸数量","卡数量","卡类","卡"])
def test_card_only_order_needs_no_auxiliary_sheet_or_history_number(header):
    headers=MINIMAL_HEADERS[:5]+[header]
    rows,_=_parse_rows("history.xlsx",workbook([row(**{header:35})],headers))
    assert len(rows)==1 and rows[0]["order_no"]==""
    assert rows[0]["line"].packaging_type=="卡纸"
    assert rows[0]["line"].required_quantity==35 and rows[0]["line"].unit=="张"


@pytest.mark.parametrize("quantity",[-1,"平卡",0,None])
def test_invalid_or_empty_card_only_demand_is_not_silently_imported(quantity):
    with pytest.raises(HTTPException):
        _parse_rows("history.xlsx",workbook([row(卡纸需求数量=quantity)],MINIMAL_HEADERS[:5]+["卡纸需求数量"]))


def test_two_card_quantity_columns_are_not_silently_reduced_to_the_first():
    with pytest.raises(HTTPException) as exc:
        _parse_rows("history.xlsx",workbook([row(卡纸需求数量=30,卡类=40)],MINIMAL_HEADERS+["卡类"]))
    assert "出现多列" in exc.value.detail


def test_inner_and_outer_keep_separate_material_and_authoritative_quantity():
    rows,_=_parse_rows("history.xlsx",workbook([row(产品订单数量=3600,内箱需求数量=1800,内箱纸质="B3B",内箱规格="10*10*10",内箱每箱个数=2,外箱需求数量=905,外箱每箱个数=4)]))
    assert len(rows)==2
    assert rows[0]["line"].paper_quality=="B3B"
    assert rows[1]["line"].required_quantity==905
    assert any("900" in warning for warning in rows[1]["row_warnings"])


def test_missing_demand_only_calculates_with_explicit_product_and_packing():
    rows,_=_parse_rows("history.xlsx",workbook([row(产品订单数量=101,外箱需求数量=None,装箱数="0/12")]))
    assert rows[0]["line"].required_quantity==9
    assert rows[0]["row_warnings"]
    with pytest.raises(HTTPException) as exc:
        _parse_rows("history.xlsx",workbook([row(外箱需求数量=None,装箱数="0/12")]))
    assert "需求数量" in exc.value.detail


@pytest.mark.parametrize("changes,fragment",[
    ({"下单日期":"9月2日"},"四位年份"),
    ({"外箱需求数量":-2},"非负"),
    ({"装箱数":"0/0"},"装箱数为0"),
    ({"装箱数":"0/12","外箱每箱个数":6},"不一致"),
    ({"计划交期":None},"计划交期"),
])
def test_ambiguous_or_invalid_values_block(changes,fragment):
    with pytest.raises(HTTPException) as exc:
        _parse_rows("history.xlsx",workbook([row(**changes)]))
    assert fragment in exc.value.detail


def test_duplicate_rows_rejected_and_independent_history_numbers_preserved():
    with pytest.raises(HTTPException) as exc:
        _parse_rows("history.xlsx",workbook([row(),row()]))
    assert "完全重复" in exc.value.detail
    parsed,_=_parse_rows("history.xlsx",workbook([row(历史订单号="H1"),row(历史订单号="H2")]))
    assert len(_group_rows(parsed))==2


def test_supplemental_cards_attach_to_only_named_main_order():
    content=workbook([row(历史订单号="H1")],extras=[["H1","平卡",200,"B3B","10*20","张"]])
    rows,_=_parse_rows("history.xlsx",content)
    assert len(_group_rows(rows))==1
    assert rows[1]["line"].packaging_type=="平卡"
    assert rows[1]["line"].required_quantity==200
    assert rows[1]["line"].usage_quantity is None
    with pytest.raises(HTTPException) as exc:
        _parse_rows("history.xlsx",workbook([row()],extras=[["H1","平卡",200,"B3B","10*20","张"]]))
    assert "未匹配" in exc.value.detail


def test_generic_material_is_not_copied_across_different_papers():
    headers=HEADERS+["纸质"]
    with pytest.raises(HTTPException) as exc:
        _parse_rows("history.xlsx",workbook([row(内箱需求数量=10,纸质="A33")],headers))
    assert "分别填写" in exc.value.detail


def test_mixed_old_and_new_sheets_never_silently_drop_old_orders():
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(workbook([row()])))
    legacy=wb.create_sheet("旧格式订单")
    legacy.append(["客户名称","下单日期","合同号","货号","纸品类型","每箱个数"])
    legacy.append(["迪奇","2026-09-01","OLD-1","ITEM-1","外箱",12])
    output=BytesIO();wb.save(output)
    with pytest.raises(HTTPException) as exc:
        _parse_rows("mixed.xlsx",output.getvalue())
    assert "混合" in exc.value.detail
    assert "旧格式订单" in exc.value.detail
