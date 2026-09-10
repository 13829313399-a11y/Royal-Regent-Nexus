from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace
from importlib import import_module

import openpyxl
import pytest
from fastapi import HTTPException

from app.api import customer_order as api
from app.services import customer_order_unified as unified, customer_order_ubtech as ubtech
from app.schemas.customer_order import CustomerOrderImportPreviewOut
from test_customer_order_huaxing_unified import template, data


TABLE = [['行\n号', '物料代码', '物料名称', '规格型号', '数量', '单位', '单价', '税率', '金额', '需求日期', '备注'],
         ['1', '3090000764', '白色时装', '规格', '1500.00', 'PCS', '17.7700', '13', '26655.00', '2026/06/24', ''],
         ['2', '3090000765', '粉色时装', '规格', '1500.00', 'PCS', '17.7700', '13', '26655.00', '2026/06/24', ''],
         ['合计', None, None, None, '3000.00', '', '', '', '53310.00', '', '']]
TEXT = '优必选 UBTECH Purchase Order 日期 : 2026/06/09 订单号 : POORD062232量产 币别:人民币 ROHS'


@pytest.fixture(autouse=True)
def current_app_modules():
    # Other API suites reload app modules while switching test databases.
    global api, unified, ubtech
    api = import_module('app.api.customer_order')
    unified = import_module('app.services.customer_order_unified')
    ubtech = import_module('app.services.customer_order_ubtech')


def mock_pdf(monkeypatch, table=None, text=TEXT):
    class Document:
        pages = [SimpleNamespace(extract_text=lambda: text, extract_tables=lambda: [deepcopy(table or TABLE)])]
        def __enter__(self): return self
        def __exit__(self, *_): pass
    monkeypatch.setattr(ubtech.pdfplumber, 'open', lambda *_: Document())


def args():
    wb = template('ubtech')
    ws = wb['ITEM表']
    ws['AD3'] = '备注'
    for r, sku in [(4, '3090000764'), (5, '3090000765')]:
        ws.cell(r, 3, '历史生产号')
        ws.cell(r, 6, 'POORD062232')
        ws.cell(r, 7, '优必选')
        ws.cell(r, 8, sku)
        ws.cell(r, 9, '历史中文名称' + sku)
        ws.cell(r, 11, 1500)
        ws.cell(r, 12, 40)
        ws.cell(r, 18, '旧箱唛POORD059826')
        ws.cell(r, 22, '是')
        wb['接单表'].cell(r, 5, 'POORD062232')
        wb['接单表'].cell(r, 7, sku)
        wb['接单表'].cell(r, 14, 20.9)
    wb['接单表']['P1'] = 7.6
    wb['正单评审表']['P1'] = "='接单表'!P1"
    return dict(customer_code='ubtech', factory_id='huakang-d', received_date='2026-09-09',
                po_files=[('PO.pdf', b'pdf')], schedule_file_name='排期.xlsx', schedule_content=data(wb))


def test_real_layout_values_and_source_priority(monkeypatch):
    mock_pdf(monkeypatch)
    preview = unified.create_unified_customer_preview(**args())
    assert preview['factory_id'] == 'huakang-d'
    assert preview['target_template'] == ubtech.TARGET_TEMPLATE
    assert len(preview['rows']) == 2
    for row in preview['rows']:
        assert row['requested_ship_date'] == '2026-06-24'
        assert row['received_date'] == '2026-09-09'
        assert row['unit_price_hkd'] == row['production_no'] == row['contract_no'] == row['carton_mark'] == ''
        assert row['reference_no'] == row['po_no'] == 'POORD062232'
        assert row['carton_count'] == '38'
        assert row['standard'] == 'ROHS'
        assert any(i['code'] == 'duplicate_existing_order' for i in row['issues'])
        date_issue = next(i for i in row['issues'] if i['field'] == 'requested_ship_date')
        assert date_issue['severity'] == 'warning' and date_issue['can_edit']
    CustomerOrderImportPreviewOut.model_validate(dict(preview, preview_fingerprint=api._preview_fingerprint(preview, '2026-09-09')))


def test_export_separates_prices_and_applies_manual_date_and_factory_price(monkeypatch):
    mock_pdf(monkeypatch)
    params = args()
    preview = unified.create_unified_customer_preview(**params)
    keys = {i['skip_key'] for row in preview['rows'] for i in row['issues'] if i['can_skip']}
    edits = [dict(row_id='ubtech-1', issue_key='ubtech-1|ubtech_ship_date_review|requested_ship_date', field='requested_ship_date', value='2026-09-30'),
             dict(row_id='ubtech-1', issue_key='ubtech-1|ubtech_factory_price_manual|unit_price_hkd', field='unit_price_hkd', value='18.5')]
    output, _, written = unified.export_unified_customer_schedule(**params, skipped_issue_keys=keys, manual_overrides=edits)
    wb = openpyxl.load_workbook(BytesIO(output))
    ws = wb['ITEM表']
    assert ws['Z6'].value.date().isoformat() == '2026-09-30'
    assert ws['C6'].value is None and ws['R6'].value is None and ws['V6'].value is None
    assert ws['M6'].value == '=CEILING(K6/L6,1)'
    assert ws['AE6'].value == 17.77 and ws['AF6'].value == '=AE6/0.85'
    assert ws['AG6'].value == '=ROUND(K6*AF6,2)'
    assert ws['AH6'].value == '=AF6/7.75'
    assert wb['接单表']['N6'].value == 18.5 and wb['接单表']['N7'].value is None
    assert wb['正单评审表']['N4'].value == '=IF(LEN(\'接单表\'!N6)=0,"",\'接单表\'!N6)'
    assert wb['接单表']['P1'].value == 7.6
    assert wb['接单表']['N4'].value == 20.9
    assert ws['R4'].value == '旧箱唛POORD059826'
    assert written['rows'][0]['carton_count'] == '38'
    assert '订单单价RMB' not in [c.value for c in openpyxl.load_workbook(BytesIO(params['schedule_content']))['ITEM表'][3]]


@pytest.mark.parametrize('column,value', [(4, '1499'), (5, 'KGM'), (6, '20'), (9, 'invalid'), (1, '')])
def test_malformed_po_cannot_silently_skip_or_change_amount(monkeypatch, column, value):
    table = deepcopy(TABLE)
    table[1][column] = value
    mock_pdf(monkeypatch, table)
    with pytest.raises(unified.CustomerOrderUnifiedError):
        ubtech.parse_po(b'pdf', 'invalid.pdf')


def test_wrong_customer_and_total_are_rejected(monkeypatch):
    mock_pdf(monkeypatch, text='Other Customer Purchase Order')
    with pytest.raises(unified.CustomerOrderUnifiedError): ubtech.parse_po(b'pdf', 'other.pdf')
    table = deepcopy(TABLE)
    table[-1][8] = '99999'
    mock_pdf(monkeypatch, table)
    with pytest.raises(unified.CustomerOrderUnifiedError, match='总数量或总金额'): ubtech.parse_po(b'pdf', 'total.pdf')


@pytest.mark.parametrize('factory', ['huaxing', 'huadeng', 'huakang-a', 'huakang-c'])
def test_customer_factory_isolation(monkeypatch, factory):
    mock_pdf(monkeypatch)
    with pytest.raises(HTTPException): api._ensure_customer_factory('ubtech', factory)
    with pytest.raises(unified.CustomerOrderUnifiedError, match='不属于当前厂区'):
        unified.create_unified_customer_preview(**dict(args(), factory_id=factory))
    api._ensure_customer_factory('ubtech', 'huakang-d')
    assert api._get_mapped_customer_spec('ubtech', 'huakang-d') == ubtech.SPEC


def test_unknown_item_does_not_inherit_fixed_pack_or_old_price(monkeypatch):
    table = deepcopy(TABLE)
    table[1][1] = '9999999999'
    mock_pdf(monkeypatch, table)
    row = unified.create_unified_customer_preview(**args())['rows'][0]
    assert row['units_per_carton'] == row['carton_count'] == row['unit_price_hkd'] == ''


def test_conflicting_pack_is_not_guessed(monkeypatch):
    mock_pdf(monkeypatch)
    params = args()
    wb = openpyxl.load_workbook(BytesIO(params['schedule_content']))
    ws = wb['ITEM表']
    ws['F7'], ws['G7'], ws['H7'], ws['L7'] = 'OTHERPO', '优必选', '3090000764', 48
    params['schedule_content'] = data(wb)
    row = unified.create_unified_customer_preview(**params)['rows'][0]
    assert row['units_per_carton'] == row['carton_count'] == ''


def test_extension_headers_can_move_but_cannot_be_ambiguous():
    wb = template('ubtech')
    ws = wb['ITEM表']
    for c, header in enumerate(reversed(ubtech.EXTENSIONS), 32): ws.cell(3, c, header)
    ubtech.write_extensions(ws, 4, {'_ubtech_price_rmb':'17.77'})
    assert ws['AL4'].value == 17.77
    assert ws['AK4'].value == '=AL4/0.85'
    assert ws['AM3'].value is None
    ws['AN3'] = '订单单价RMB'
    with pytest.raises(unified.CustomerOrderUnifiedError, match='重复'):
        ubtech.write_extensions(ws, 5, {'_ubtech_price_rmb':'17.77'})
