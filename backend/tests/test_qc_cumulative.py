"""Real SQLite, source/photo stores, native template and OOXML regression checks."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from app.services import qc_cumulative as c, qc_inspection_sheet as sheet, qc_workbook_xml as writer
from test_qc_inspection_sheet import source, save, photo, pushed_candidates, old_source_candidates
from test_customer_order_ledger import engine,db,ingest
from test_qc_workflow import qc_user

@pytest.fixture(autouse=True)
def isolated(monkeypatch,tmp_path):monkeypatch.setattr(sheet.settings,'qc_department_data_dir',str(tmp_path/'companies'))

SCOPE={'dateFrom':'2026-01-01','dateTo':'2026-12-31','search':'','result':''}
def seed(count=3):
    sources=[]
    for i in range(count):
        row=source('new-'+str(i),'000TEST-'+str(i));row.update(customer=['ZURU','BUZZBEE','Disney'][i%3],country='USA',dateFrom='2026-10-29' if i%2==0 else '2026-11-03',dateTo='2026-11-03',remark='拉杆松，客户要求周五补验')
        sources.append(row)
    return sheet.sync('huaxing',pushed_candidates(sources))['items']
def request(items,version=0,key=None,scope=None):return {'requestKey':key or uuid4().hex,'expectedVersion':version,'scope':scope or SCOPE,'revisions':{r['id']:r['revision'] for r in items}}
def run(items,version=0,**kwargs):return c.operate('huaxing','tester',request(items,version,**kwargs),items)
def book():return load_workbook(BytesIO(c.download('huaxing')[0]))
def latest():
    with c.store() as conn:
        h=conn.execute('SELECT * FROM heads').fetchone();v=conn.execute('SELECT * FROM versions WHERE id=?',(h['current_id'],)).fetchone()
        return dict(h),json.loads(v['snapshot']),json.loads(v['metadata'])

def test_template_history_and_mapping_with_photos_and_blank_results():
    items=seed();items[0],_=photo(items[0]);before=c.download('huaxing')[0]
    result=run(items);assert result['count']==3 and result['version']==1
    data=c.download('huaxing')[0];w=load_workbook(BytesIO(data));base=load_workbook(BytesIO(before))
    assert w.sheetnames[0]=='验货明细' and '累计验货详情' not in w.sheetnames
    _,snap,meta=latest()
    assert w['验货明细'].freeze_panes=='D2' and len(w['验货明细']._images)==0
    for row in items:
        binding=meta['bindings'][row['id']];s=w[binding['sheet']];r=binding['row'];second=row['customer']=='BUZZBEE'
        assert s.cell(r,11 if second else 10).value==row['remark']
        assert s.cell(r,13).value in ('',None)
        assert s.cell(r,8 if second else 9).value in ('',None)
        assert s.cell(r,2 if second else 3).value=='USA'
        assert s.cell(r,1 if second else 2).hyperlink is None
        summary=meta['summaryRows'][row['id']]
        assert w[summary['sheet']].cell(summary['row'],3).value==row['poNo']
    # Every pre-existing value, formula, hidden state, merge and print margin survives.
    for name in base.sheetnames:
        a,b=base[name],w[name];assert a.sheet_state==b.sheet_state and set(a.merged_cells.ranges)<=set(b.merged_cells.ranges) and a.page_margins==b.page_margins
        for row in a:
            for cell in row:
                if cell.value is not None:
                    v1,v2=cell.value,b[cell.coordinate].value
                    if hasattr(v1,'text'):assert v1.text==v2.text
                    else:assert v1==v2,(name,cell.coordinate)
    with ZipFile(BytesIO(before)) as original,ZipFile(BytesIO(data)) as updated:
        pkg=writer.Package(before)
        for name in writer.PROTECTED:assert original.read(pkg.path(name))==updated.read(pkg.path(name))

def test_noop_idempotency_edit_undo_and_source_unchanged():
    items=seed(1);initial=c.download('huaxing')[0];body=request(items)
    one=c.operate('huaxing','tester',body,items)
    assert c.operate('huaxing','tester',body,items)==one
    data1=c.download('huaxing')[0]
    assert run(items,1)['status']=='unchanged'
    assert c.download('huaxing')[0]==data1
    items[0]=save(items[0],'FAIL','完整原文\n客户要求补验')
    assert run(items,1)['version']==2
    undo={'requestKey':uuid4().hex,'expectedVersion':2}
    assert c.operate('huaxing','tester',undo,[],'undo')['version']==3
    assert c.download('huaxing')[0]==data1
    with sheet.store() as conn:assert sheet.get(conn,'huaxing',items[0]['id'])['result']=='FAIL'
    with pytest.raises(HTTPException):c.operate('huaxing','tester',{'requestKey':uuid4().hex,'expectedVersion':3},[],'undo')
    assert run(items,3)['version']==4
    assert c.job('huaxing','tester',body['requestKey'])['status']=='succeeded'
    assert c.job('huaxing','other',body['requestKey'])['status']=='unknown'

def test_empty_and_all_reference_batches_preserve_undo():
    items=seed(1);run(items);before=c.download('huaxing')[0]
    demo={**items[0],'testData':True};assert run([demo],1)['status']=='empty'
    assert c.download('huaxing')[0]==before and latest()[0]['undo_id'] is not None
    assert run([{**items[0],'referenceData':True}],1)['status']=='empty'
    assert c.state('huaxing',[demo],SCOPE)['excluded']==1

def test_missing_dates_failed_batch_and_stale_revision_atomic(monkeypatch):
    items=seed(2);before=c.download('huaxing')[0]
    broken=deepcopy(items);broken[1].update(actualDate='',dateFrom='',dateTo='')
    with pytest.raises(HTTPException,match='缺少'):run(broken)
    assert c.download('huaxing')[0]==before
    saved=save(items[0],'PASS','new')
    with pytest.raises(HTTPException) as e:run(items)
    assert e.value.status_code==409 and c.download('huaxing')[0]==before
    items[0]=saved
    def fail(*args):raise RuntimeError('injected disk/writer failure')
    monkeypatch.setattr(writer,'build',fail)
    with pytest.raises(HTTPException):run(items)
    assert c.download('huaxing')[0]==before and latest()[0]['undo_id'] is None

def test_actual_date_migration_single_identity_and_crossweek_plan():
    items=seed(1);items[0].update(dateFrom='2026-11-01',dateTo='2026-11-03')
    # Persist the same source base as the UI would after authoritative source sync.
    run(items);_,snapshot,_=latest();assert snapshot[items[0]['id']]['week']=='2026-10-26'
    items[0]['actualDate']='2026-11-03';run(items,1)
    _,snapshot,meta=latest();assert len(snapshot)==1 and snapshot[items[0]['id']]['week']=='2026-11-02'
    w=book();assert w.sheetnames[0]=='验货明细' and w['验货明细'].max_row==2
    assert sum(cell.value==items[0]['poNo'] for row in w['ZURU行验货'] for cell in row)==1

def test_long_verbatim_text_photo_removal_no_broken_relationships():
    items=seed(1);items[0]=save(items[0],'PASS','中文字\n'*3000);items[0],_=photo(items[0]);run(items)
    w=book();_,_,meta=latest();binding=meta['bindings'][items[0]['id']]
    assert w[binding['sheet']].cell(binding['row'],10).value==items[0]['remark']
    assert max(r.height or 0 for r in w[binding['sheet']].row_dimensions.values())<=409.5
    assert w['验货明细']['O2'].value==items[0]['remark']
    items[0]=sheet.remove_photo('huaxing',items[0]['id'],{'revision':items[0]['revision'],'photoId':items[0]['photos'][0]['id']},{'id':'tester'})
    items[0]=save(items[0],'FAIL','短备注');run(items,1);w=book();assert len(w['验货明细']._images)==0
    with ZipFile(BytesIO(c.download('huaxing')[0])) as z:assert not any('qc-photo-' in p for p in z.namelist())

def test_parallel_import_and_undo_cas():
    items=seed(1)
    def call():
        try:return run(items)
        except HTTPException as e:return e.status_code
    with ThreadPoolExecutor(2) as pool:results=list(pool.map(lambda _:call(),range(2)))
    assert sum(isinstance(r,dict) for r in results)==1 and 409 in results
    with pytest.raises(HTTPException) as e:c.operate('huaxing','tester',{'requestKey':uuid4().hex,'expectedVersion':0},[],'undo')
    assert e.value.status_code==409

def test_template_factory_isolation_and_exact_quantity_validation():
    with pytest.raises(HTTPException) as e:c.download('huadeng')
    assert e.value.status_code==503
    items=seed(1);items[0]['cartonCount']='5'
    with pytest.raises(HTTPException,match='不一致'):run(items)
    assert latest()[0]['revision']==0

def test_source_to_workflow_to_legacy_keeps_identity_and_reinspection_separate():
    src=source();first=sheet.sync('huaxing',old_source_candidates([src]))['items'][0];first=save(first,'PASS','已检')
    task={'id':'workflow-a','sourceId':src['id'],'source':src,'batch':'batch-1','execution':{'result':'','remark':'','actualDate':''},'attachments':[]}
    rows=sheet.sync('huaxing',sheet.candidates('huaxing',[],[task],[src]))['items'];assert len(rows)==1 and rows[0]['id']==first['id'] and rows[0]['remark']=='已检'
    linked={**task,'legacyId':100,'legacyOriginal':{'id':100,'client':'客户甲','orderNo':src['poNo'],'productNo':src['productNo'],'qty':24}}
    rows=sheet.sync('huaxing',sheet.candidates('huaxing',[],[linked],[src]))['items'];assert rows[0]['id']==first['id']
    again={**task,'id':'workflow-b','batch':'batch-2'}
    rows=sheet.sync('huaxing',sheet.candidates('huaxing',[],[linked,again],[src]))['items'];assert len(rows)==2 and len({r['id'] for r in rows})==2

def test_unique_historical_binding_preserves_formula_and_migrates_without_duplicate():
    items=seed(1);p=writer.Package(c.template('huaxing'));first=p.rows('ZURU行验货')[3];vals=p.values(first)
    # Pick the first real legacy row (date group 2026-01-06).
    row=next(r for r in p.rows('ZURU行验货') if p.values(r).get(2)==4500188353)
    values=p.values(row);items[0].update(poNo='4500188353',productNo=values[4],quantity=values[6],unitsPerCarton=values[7],cartonCount=values[8],dateFrom='2026-01-06',dateTo='2026-01-06')
    run(items);assert latest()[2]['bindings'][items[0]['id']].get('original')
    items[0]['actualDate']='2026-11-03';run(items,1)
    assert sum(str(cell.value)=='4500188353' for r in book()['ZURU行验货'] for cell in r)==1

def test_repeat_same_content_after_save_is_noop_and_pending_zero():
    items=seed(1);run(items);before=c.download('huaxing')[0]
    items[0]=save(items[0],items[0]['result'],items[0]['remark'])
    assert run(items,1)['status']=='unchanged'
    assert c.download('huaxing')[0]==before and c.state('huaxing',items,SCOPE)['pending']==0

def test_actual_http_import_download_undo_and_permissions(db,monkeypatch):
    from dataclasses import replace
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    import httpx
    from app.api import qc_department as gateway
    from app.db import get_db
    from app.services import auth
    from app.services.auth import get_current_user
    monkeypatch.setattr(auth,'add_auth_audit',lambda *a,**k:None)
    async def upstream(*args):return httpx.Response(200,json={'records':[]})
    monkeypatch.setattr(gateway,'upstream',upstream)
    line=ingest(db);line.data={**line.data,'planned_inspection_date':'2026-10-29'};db.commit()
    app=FastAPI();app.include_router(gateway.router);current=[qc_user()]
    app.dependency_overrides[get_current_user]=lambda:current[0];app.dependency_overrides[get_db]=lambda:db
    with TestClient(app) as client:
        base=client.get('/api/qc-department/huaxing/api/health').json()['frame_url']+'api/inspection-sheet/'
        assert client.get(base+'entries').json()['items']==[]
        pushed=client.post(base.replace('inspection-sheet/','workflow/tasks'),json={'requestKey':uuid4().hex,'items':[{'sourceId':line.id,'sourceRevision':line.revision,'method':'internal','plannedDate':'2026-10-29','plannedQuantity':str(line.quantity),'batch':'首验'}]})
        assert pushed.status_code==200,pushed.text
        items=client.get(base+'entries').json()['items'];assert len(items)==1
        before=client.get(base+'workbook/download').content
        body=request(items)
        result=client.post(base+'workbook/import',json=body);assert result.status_code==200,result.text
        assert len(load_workbook(BytesIO(client.get(base+'workbook/download').content)).sheetnames)==17
        assert client.get(base+'workbook/jobs/'+body['requestKey']).json()['status']=='succeeded'
        assert client.post(base+'workbook/undo',json={'requestKey':uuid4().hex,'expectedVersion':1}).status_code==200
        assert client.get(base+'workbook/download').content==before
        assert client.post(base+'workbook/import',json=request(items,2),headers={'Sec-Fetch-Site':'cross-site'}).status_code==403
        assert client.get(base+'workbook/download?company=hy-huadeng').status_code==403
        row=items[0];with_photo,_=photo(row)
        mapped_body={'revision':row['revision'],'result':'FAIL','remark':'mapped HTTP save'}
        saved=client.post(base+'workbook/entries/'+row['id']+'/save',json=mapped_body)
        assert saved.status_code==200 and saved.json()['result']=='FAIL' and len(saved.json()['photos'])==1
        assert client.post(base+'workbook/entries/'+row['id']+'/save',json=mapped_body).status_code==409
        response=client.head(base+'workbook/download');assert response.status_code==200 and not response.content
        assert int(response.headers['content-length'])==len(before)
        replacement_body=replacement(fresh_base(),version=2)
        replaced=client.post(base+'workbook/replace',json=replacement_body)
        assert replaced.status_code==200,replaced.text
        archive_id=replaced.json()['archiveId']
        assert client.get(base+'workbook/archives').json()['items'][0]['id']==archive_id
        assert client.get(base+'workbook/archives/'+str(archive_id)).content==before
        assert client.get(base+'workbook/archives/'+str(archive_id)+'?company=hy-huadeng').status_code==403
        assert client.post(base+'workbook/replace',json=replacement_body,headers={'Sec-Fetch-Site':'cross-site'}).status_code==403
        assert client.get(base+'entries').json()['items'][0]['result']=='FAIL'
        user=current[0];permissions=frozenset({'qc_inspection:read'})
        current[0]=replace(user,permissions=permissions,grants=tuple(replace(g,permissions=permissions) for g in user.grants))
        assert client.get(base+'workbook/archives').status_code==403
        assert client.get(base+'workbook/archives/'+str(archive_id)).status_code==403
        assert client.post(base+'workbook/replace',json=replacement_body).status_code==403
        assert client.get(base+'workbook/download').status_code==403
        assert client.post(base+'workbook/import',json=request(items,2)).status_code==403
        assert client.post(base+'workbook/entries/'+row['id']+'/save',json=mapped_body).status_code==403


def test_all_actual_judgments_use_one_snapshot_and_never_workflow_progress(monkeypatch):
    from app.services import qc_workflow as wf
    supported=['','PASS','FAIL','REJ','COND','HOLD']
    assert {'','PASS','REJ','COND','HOLD'}==wf.RESULTS
    candidates=[]
    for i,result in enumerate(supported):
        src=source('judgment-'+str(i),'000QA-CUM-'+str(i).zfill(4));src['customer']='ZURU'
        task={'id':'judgment-'+str(i),'sourceId':src['id'],'source':src,'batch':'1','execution':{'result':result,'remark':'拉杆松，客户要求周五补验'},
            'attachments':[{'id':'expired','mime':'image/jpeg','url':'https://invalid.example/expired'}]}
        candidates.extend(sheet.candidates('huaxing',[],[task],[src]))
    monkeypatch.setattr(wf,'read_file',lambda *args:pytest.fail('cumulative synchronization read photo bytes'))
    items=sheet.sync('huaxing',candidates,False)['items'];assert len(items)==6
    for item in items:item['completion']='未推送'
    run(items);w=book();_,snapshot,meta=latest()
    for item,expected in zip(items,supported):
        identity=item['id'];binding=meta['bindings'][identity];summary=meta['summaryRows'][identity]
        assert snapshot[identity]['result']==expected
        assert (w[binding['sheet']].cell(binding['row'],9).value or '')==expected
        assert (w[summary['sheet']].cell(summary['row'],14).value or '')==expected
    reject=items[3];reject=save(reject,reject['result'],'仅改备注');assert reject['cumulativeResult']=='REJ'
    hold=items[5];hold=save(hold,'PASS','本人已判定');assert hold['cumulativeResult']=='PASS'
    blank=items[0];blank=save(blank,'FAIL','新结论');run([blank],1)
    blank=save(blank,'PASS','复核');run([blank],2)
    assert len(latest()[1])==6
    _,_,meta=latest();r=meta['summaryRows'][blank['id']];assert book()[r['sheet']].cell(r['row'],14).value=='PASS'


def test_detail_reuses_original_export_and_customer_mapping_without_subtotals():
    from datetime import datetime
    from copy import copy
    sources=[]
    for client in [*writer.CLIENTS,'Disney']:
        for suffix in ('003-B','001-A'):
            src=source(client+suffix,'000 '+suffix+'-01234567890123456789-X')
            src.update(customer=client,dateFrom='2026-10-29',dateTo='2026-10-29',productNo='000 SKU '+suffix,remark='中文长备注\n'+('需要复查外箱标识及装配间隙。'*20),country='USA')
            sources.append(src)
    items=sheet.sync('huaxing',pushed_candidates(sources),False)['items']
    run(items[:11]);run(items[11:],1);assert book()['验货明细'].max_row==len(items[11:])+1
    run(items,2);w=book();_,snap,meta=latest();ordered=[snap[i] for i in meta['detailIds']]
    reference=load_workbook(BytesIO(sheet.export_records(ordered,include_photos=False))).active
    actual=w.worksheets[0];assert actual.title=='验货明细' and actual.max_column==15 and actual.max_row==25
    assert actual.freeze_panes==reference.freeze_panes and actual.auto_filter.ref==reference.auto_filter.ref
    assert actual.print_title_rows==reference.print_title_rows and actual.page_setup==reference.page_setup
    for row in reference:
        for expected in row:
            got=actual[expected.coordinate]
            assert got.value==expected.value and got.number_format==expected.number_format
            assert copy(got.font)==copy(expected.font) and copy(got.fill)==copy(expected.fill) and copy(got.border)==copy(expected.border) and copy(got.alignment)==copy(expected.alignment)
    for col,dim in reference.column_dimensions.items():assert actual.column_dimensions[col].width==dim.width
    assert [actual.row_dimensions[i].height for i in range(1,26)]==[reference.row_dimensions[i].height for i in range(1,26)]
    for item in items:
        loc=meta['bindings'][item['id']];name=loc['sheet'];row=loc['row'];second=name in writer.SECOND
        assert w[name].cell(row,1 if second else 2).value==item['poNo']
        assert w[name].cell(row,3 if second else 4).value==item['productNo']
        assert w[name].cell(row,11 if second else 10).value==item['remark']
    assert all('照片' not in str(c.value) for c in actual[1]) and not actual._images
    assert not any(c.value in ('客户小计','当日小计') for row in actual for c in row)


def test_photo_only_changes_are_noop_and_mapped_save_ignores_only_audited_photos(monkeypatch):
    from contextlib import contextmanager
    import sqlite3
    items=seed(1);run(items);before=c.download('huaxing')[0];head=latest()[0];old_revision=items[0]['revision']
    changed,_=photo(items[0]);items[0]=changed
    assert run(items,1)['status']=='unchanged';assert latest()[0]==head and c.download('huaxing')[0]==before
    # A stale client revision is safe only when every intervening revision is a photo action.
    saved=sheet.save('huaxing',changed['id'],{'revision':old_revision,'result':'FAIL','remark':'照片失败不阻塞文字'}, {'id':'tester'},True)
    with pytest.raises(HTTPException):sheet.save('huaxing',saved['id'],{'revision':old_revision,'result':'PASS','remark':'stale'}, {'id':'tester'},True)
    original_store=sheet.store
    @contextmanager
    def no_photos(write=False):
        with original_store(write) as conn:
            conn.set_authorizer(lambda action,table,*rest:sqlite3.SQLITE_DENY if action==sqlite3.SQLITE_READ and table=='photos' else sqlite3.SQLITE_OK)
            yield conn
    monkeypatch.setattr(sheet,'store',no_photos)
    saved['photos']=[{'id':'broken','url':'https://invalid.example/expired'}]
    assert c.state('huaxing',[saved],SCOPE)['pending']==1
    assert run([saved],1)['status']=='imported'
    assert latest()[1][saved['id']].get('photoDigests') is None


def make_legacy_version():
    items=seed(1);items[0]['poNo']='000QA-CUM-0000';run(items)
    items[0]=save(items[0],'FAIL','问题原文');items[0]['poNo']='000QA-CUM-0000';run(items,1)
    with c.store(True) as conn:
        head=conn.execute('SELECT * FROM heads').fetchone()
        for identity in (head['current_id'],head['undo_id']):
            row=conn.execute('SELECT * FROM versions WHERE id=?',(identity,)).fetchone();snapshot=json.loads(row['snapshot'])
            content,meta=writer._build(c.template('huaxing'),{},snapshot,{},legacy=True);meta['formatVersion']=3
            # Construct the actual prior layout from saved snapshots, not a damaged current-format page.
            conn.execute('UPDATE versions SET content=?,metadata=? WHERE id=?',(content,json.dumps(meta),identity))
    return items


def test_persisted_repair_and_repair_rollback_use_snapshots_and_correct_undo(monkeypatch):
    items=make_legacy_version();before=c.download('huaxing')[0];oldhead=latest()[0]
    with c.store() as conn:oldrows=[tuple(r) for r in conn.execute('SELECT * FROM versions')]
    original_store=sheet.store
    monkeypatch.setattr(sheet,'store',lambda *args:pytest.fail('repair must not read live inspection records'))
    result=c.repair('huaxing','repair-test',2);assert result['status']=='repaired'
    _,snap,meta=latest();loc=meta['summaryRows'][items[0]['id']];assert book()[loc['sheet']].cell(loc['row'],14).value=='FAIL'
    with c.store() as conn:assert [tuple(r) for r in conn.execute('SELECT * FROM versions ORDER BY id')][:len(oldrows)]==oldrows
    c.rollback_repair('huaxing',result['repairId'],3);assert c.download('huaxing')[0]==before
    assert latest()[0]['undo_id']==oldhead['undo_id']
    with pytest.raises(HTTPException):c.rollback_repair('huaxing',result['repairId'],4)
    c.repair('huaxing','repair-test',4)
    c.operate('huaxing','tester',{'requestKey':uuid4().hex,'expectedVersion':5},[],'undo')
    _,snap,meta=latest();loc=meta['summaryRows'][items[0]['id']]
    assert (book()[loc['sheet']].cell(loc['row'],14).value or '')==''
    assert snap[items[0]['id']]['result']==''
    monkeypatch.setattr(sheet,'store',original_store)
    assert run(items,6)['status']=='imported'


def test_version_insert_failure_preserves_file_snapshot_and_undo(monkeypatch):
    from contextlib import contextmanager
    import sqlite3
    items=seed(1);run(items);items[0]=save(items[0],'FAIL','changed')
    before=c.download('huaxing')[0];head,snap,meta=latest();original=c.store
    @contextmanager
    def denied(write=False):
        with original(write) as conn:
            conn.set_authorizer(lambda action,table,*rest:sqlite3.SQLITE_DENY if action==sqlite3.SQLITE_INSERT and table=='versions' else sqlite3.SQLITE_OK)
            yield conn
    monkeypatch.setattr(c,'store',denied)
    with pytest.raises(HTTPException):run(items,1)
    assert c.download('huaxing')[0]==before and latest()==(head,snap,meta)


def test_repair_failure_is_atomic_and_download_does_not_update_store(monkeypatch):
    import sqlite3
    items=make_legacy_version();before=c.download('huaxing')[0];original=c.writer.build;calls=[];head=latest()[0]
    def fail_second(*args):
        calls.append(1)
        if len(calls)==2:raise RuntimeError('second repaired version failed')
        return original(*args)
    monkeypatch.setattr(c.writer,'build',fail_second)
    with pytest.raises(RuntimeError):c.repair('huaxing','tester',2)
    assert latest()[0]==head and c.download('huaxing')[0]==before
    monkeypatch.setattr(c,'store',lambda *args:pytest.fail('download tried mutable store'))
    assert c.download('huaxing')[0]==before


def test_original_picture_with_old_generated_prefix_is_preserved(monkeypatch):
    from test_qc_workflow import png
    pkg=writer.Package(c.template('huaxing'));name='新明自检';sheet_path=pkg.path(name)
    image='xl/media/qc-photo-historical.png';drawing='xl/drawings/qc-photo-historical.xml'
    pkg.parts[image]=png()
    pkg.parts[drawing]=('''<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><xdr:oneCellAnchor><xdr:from><xdr:col>0</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>0</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:ext cx="914400" cy="914400"/><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="1" name="Original historical image"/><xdr:cNvPicPr/></xdr:nvPicPr><xdr:blipFill><a:blip r:embed="historical"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill><xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr></xdr:pic><xdr:clientData/></xdr:oneCellAnchor></xdr:wsDr>''').encode()
    pkg.parts['xl/drawings/_rels/qc-photo-historical.xml.rels']=(f'<Relationships xmlns="{writer.P}"><Relationship Id="historical" Type="{writer.R}/image" Target="../media/qc-photo-historical.png"/></Relationships>').encode()
    rel_path='xl/worksheets/_rels/'+sheet_path.rsplit('/',1)[1]+'.rels'
    old_rels=writer.E.fromstring(pkg.parts[rel_path]) if rel_path in pkg.parts else writer.E.Element('{'+writer.P+'}Relationships')
    writer.E.SubElement(old_rels,'{'+writer.P+'}Relationship',{'Id':'historicalPicture','Type':writer.R+'/drawing','Target':'../drawings/qc-photo-historical.xml'})
    pkg.parts[rel_path]=writer.xml(old_rels);writer.E.SubElement(pkg.sheet(name),writer.tag('drawing'),{'{'+writer.R+'}id':'historicalPicture'})
    writer.E.SubElement(pkg.types,'{'+writer.C+'}Override',{'PartName':'/'+drawing,'ContentType':'application/vnd.openxmlformats-officedocument.drawing+xml'})
    if not any(t.get('Extension')=='png' for t in pkg.types):writer.E.SubElement(pkg.types,'{'+writer.C+'}Default',{'Extension':'png','ContentType':'image/png'})
    original=pkg.finish();monkeypatch.setattr(c,'template',lambda factory:original)
    items=seed(1);run(items);items[0]=save(items[0],'FAIL','changed');run(items,1)
    updated=c.download('huaxing')[0]
    with ZipFile(BytesIO(original)) as before,ZipFile(BytesIO(updated)) as after:
        for part in (image,drawing,rel_path,sheet_path):assert before.read(part)==after.read(part)
    w=load_workbook(BytesIO(updated));assert len(w[name]._images)==1 and not w['验货明细']._images


def test_unavailable_historical_photo_does_not_hide_text_or_block_import(monkeypatch):
    from app.services import qc_workflow as wf
    src=source('unavailable-photo');src['customer']='ZURU'
    task={'id':'historical-photo-task','sourceId':src['id'],'source':src,'batch':'1','execution':{'result':'HOLD','remark':'等待补验'},
          'attachments':[{'id':'missing-photo','mime':'image/jpeg','name':'old.jpg'}]}
    rows=sheet.candidates('huaxing',[],[task],[src])
    def missing(*args):raise HTTPException(503,'attachment temporarily unavailable')
    monkeypatch.setattr(wf,'read_file',missing)
    items=sheet.sync('huaxing',rows)['items'];assert items[0]['photoWarning'] and items[0]['cumulativeResult']=='HOLD'
    assert run(items)['status']=='imported' and latest()[1][items[0]['id']]['result']=='HOLD'
    assert 'photoWarning' not in latest()[1][items[0]['id']]
    # The next normal refresh can recover the original photograph; it still does not create an Excel version.
    from test_qc_workflow import png
    monkeypatch.setattr(wf,'read_file',lambda *args:{'content':png()})
    restored=sheet.sync('huaxing',rows)['items'];assert len(restored[0]['photos'])==1 and not restored[0].get('photoWarning')
    assert run(restored,1)['status']=='unchanged'


def replacement(content,name='2027.xlsx',version=0):
    import base64
    return {'requestKey':uuid4().hex,'expectedVersion':version,'filename':name,'content':base64.b64encode(content).decode()}


def fresh_base():
    p=writer.Package(c.template('huaxing'))
    # Retain the provided template, plus a clearly independent annual marker.
    p.row('新明自检',10000,['新年度原始标记'],{},20)
    return p.finish()


def test_new_base_archives_exact_bytes_then_accumulates_pending_and_undo():
    rows=seed(1);rows[0]=save(rows[0],'PENDING','待客户确认')
    run(rows);old=c.download('huaxing')[0];source_before=deepcopy(rows);uploaded=fresh_base()
    body=replacement(uploaded,version=1)
    result=c.operate('huaxing','tester',body,[],'replace')
    assert result['status']=='replaced' and result['version']==2 and not result['undoAvailable']
    assert c.download('huaxing')[0]==uploaded
    assert c.operate('huaxing','tester',body,[],'replace')==result
    assert len(c.archives('huaxing')['items'])==1
    assert c.archive_download('huaxing',result['archiveId'])[0]==old
    with pytest.raises(HTTPException) as wrong:c.archive_download('other-factory',result['archiveId'])
    assert wrong.value.status_code==404
    with sheet.store() as conn:assert sheet.unpack(conn,sheet.get(conn,'huaxing',rows[0]['id']))==source_before[0]
    pending=c.state('huaxing',rows,{**SCOPE,'result':'PENDING'});assert pending['count']==1
    assert c.state('huaxing',rows,{**SCOPE,'result':'PASS'})['count']==0
    assert c.state('huaxing',rows,{**SCOPE,'result':'FAIL'})['count']==0
    run(rows,2);data=c.download('huaxing')[0];workbook=load_workbook(BytesIO(data));_,snap,meta=latest()
    item=rows[0];binding=meta['bindings'][item['id']];summary=meta['summaryRows'][item['id']]
    assert workbook[binding['sheet']].cell(binding['row'],9).value=='PENDING'
    assert workbook[summary['sheet']].cell(summary['row'],14).value=='PENDING'
    assert workbook['新明自检'].cell(10000,1).value=='新年度原始标记'
    assert run(rows,3)['status']=='unchanged'
    c.operate('huaxing','tester',{'requestKey':uuid4().hex,'expectedVersion':3},[],'undo')
    assert c.download('huaxing')[0]==uploaded


def test_failed_upload_stale_upload_and_same_file_leave_head_and_archives_unchanged():
    baseline=c.download('huaxing')[0]
    for content,name in [(b'broken','broken.xlsx'),(b'broken','broken.xls')]:
        with pytest.raises(HTTPException):c.operate('huaxing','tester',replacement(content,name),[],'replace')
        assert c.download('huaxing')[0]==baseline
        assert c.archives('huaxing')['items']==[]
    with pytest.raises(HTTPException) as stale:c.operate('huaxing','tester',replacement(fresh_base(),version=99),[],'replace')
    assert stale.value.status_code==409
    assert c.archives('huaxing')['items']==[]
    result=c.operate('huaxing','tester',replacement(baseline),[],'replace');assert result['status']=='unchanged'
    assert c.download('huaxing')[0]==baseline


def test_upload_rejects_wrong_layout_and_preserves_unknown_week_sheet():
    from app.services import qc_workbook_upload as upload
    p=writer.Package(fresh_base());name='ZURU行验货';profile=p.profile(name)
    row=next(r for r in p.rows(name) if int(r.get('r'))==profile['templateRows']['header']);p.cell(row,2,'错误列',0)
    with pytest.raises(HTTPException):upload.prepare('invalid.xlsx',p.finish(),c.template('huaxing'))
    rows=seed(1);p=writer.Package(fresh_base());name='2026-10-26至11-01';p.add(name,'ZURU行验货');p.rows(name)[:]=[];p.row(name,1,['用户已有的周汇总'],{},20)
    c.operate('huaxing','tester',replacement(p.finish()),[],'replace');run(rows,1)
    w=load_workbook(BytesIO(c.download('huaxing')[0]));assert w[name]['A1'].value=='用户已有的周汇总'
    assert w.sheetnames[0]=='验货明细' and name+'（续1）' not in w.sheetnames


def test_known_export_reupload_retains_managed_identity_and_does_not_duplicate():
    rows=seed(1);run(rows);exported=c.download('huaxing')[0]
    c.operate('huaxing','tester',replacement(fresh_base(),version=1),[],'replace')
    c.operate('huaxing','tester',replacement(exported,version=2),[],'replace')
    assert run(rows,3)['status']=='unchanged'
    assert c.download('huaxing')[0]==exported


def test_blank_annual_template_can_accumulate_using_retained_styles():
    from app.services import qc_workbook_upload as upload
    p=writer.Package(fresh_base())
    for name in set(writer.TARGETS.values()):
        profile=p.profile(name)
        headers={int(row.get('r')) for row in p.rows(name) if writer.normalized(p.values(row).get(1))==writer.normalized('Contract No.' if name in writer.SECOND else '客户')}
        for row in p.rows(name):
            if int(row.get('r')) in headers:continue
            for cell in row:
                cell[:]=[];cell.attrib.pop('t',None)
    blank=p.finish();assert upload.prepare('2027.xlsx',blank,c.template('huaxing'))[0]==blank
    c.operate('huaxing','tester',replacement(blank),[],'replace')
    rows=seed(1);assert run(rows,1)['status']=='imported'
    assert load_workbook(BytesIO(c.download('huaxing')[0])).sheetnames[0]=='验货明细'


def test_real_xls_converter_preserves_template_and_archives_original(monkeypatch):
    import os,shutil
    converter=os.getenv('QC_XLS_CONVERTER') or shutil.which('x2t') or '/Applications/ONLYOFFICE.app/Contents/Resources/converter/x2t'
    if not Path(converter).is_file():pytest.skip('ONLYOFFICE x2t not installed')
    original=(c.TEMPLATES/'huaxing/original.xls').read_bytes()
    result=c.operate('huaxing','tester',replacement(original,'新年度.xls'),[],'replace')
    assert result['status']=='replaced'
    assert result['conversionReceipt']['cellsVerified']==41646
    assert result['conversionReceipt']['commentsRestored']==6
    again=c.operate('huaxing','tester',replacement(original,'新年度.xls',version=1),[],'replace')
    assert again['status']=='unchanged' and again['version']==1
    assert len(c.archives('huaxing')['items'])==1
    with c.store() as conn:assert bytes(conn.execute('SELECT content FROM base_inputs').fetchone()[0])==original
    rows=seed(1);assert run(rows,1)['status']=='imported'
    w=load_workbook(BytesIO(c.download('huaxing')[0]));assert 'ZURU行验货' in w.sheetnames


def test_archive_and_switch_roll_back_together_on_storage_failure():
    before=c.download('huaxing')[0]
    with c.store(True) as conn:
        c.initialize(conn,'huaxing')
        conn.execute("CREATE TRIGGER reject_base BEFORE INSERT ON base_inputs BEGIN SELECT RAISE(ABORT, 'simulated disk failure'); END")
    body=replacement(fresh_base())
    with pytest.raises(HTTPException) as failure:c.operate('huaxing','tester',body,[],'replace')
    assert failure.value.status_code==500
    assert c.download('huaxing')[0]==before
    assert c.archives('huaxing')['items']==[]
    assert c.job('huaxing','tester',body['requestKey'])['status']=='failed'
    assert latest()[0]['revision']==0


def test_unavailable_xls_converter_keeps_active_workbook(monkeypatch):
    monkeypatch.setenv('QC_XLS_CONVERTER','/nonexistent-qc-test/x2t')
    before=c.download('huaxing')[0];original=(c.TEMPLATES/'huaxing/original.xls').read_bytes()
    with pytest.raises(HTTPException) as failure:c.operate('huaxing','tester',replacement(original,'new.xls'),[],'replace')
    assert failure.value.status_code==503
    assert c.download('huaxing')[0]==before and c.archives('huaxing')['items']==[]


def colored_base():
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill,Font,Color,Alignment
    wb=Workbook();ws=wb.active;ws.title='ZURU行验货';ws.append(writer.HEAD1)
    for cell in ws[2]:cell.font=Font(name='Arial',size=11);cell.alignment=Alignment(wrap_text=True)
    ws['J2'].fill=PatternFill('solid',fgColor='FF000000');ws['J2'].font=Font(name='Arial',color='FFFFFFFF')
    colors=wb.create_sheet('颜色和历史样例')
    for i,(fill,font) in enumerate([(PatternFill('solid',fgColor='FF000000'),Font(color='FFFFFFFF')),(PatternFill('solid',fgColor='FFFFFF00'),Font(color='FFFF0000')),(PatternFill(),Font(color=Color(auto=True))),(PatternFill('solid',fgColor=Color(theme=4,tint=.35)),Font(color='FF000000')),(PatternFill('solid',fgColor=Color(indexed=42)),Font(color='FF000000'))],1):
        cell=colors.cell(i,1,'0000123- 测试中文');cell.fill=fill;cell.font=font
    colors['B1']='=1+2';out=BytesIO();wb.save(out);return out.getvalue()


def test_color_representations_and_three_known_roundtrips():
    from copy import copy
    from app.services import qc_workbook_upload as upload
    raw=colored_base();out,receipt=upload.prepare('合法色彩.xlsx',raw,c.template('huaxing'),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    assert out==raw
    c.operate('huaxing','tester',replacement(raw),[],'replace');items=seed(1);items[0]=save(items[0],'PENDING','测试中文备注');run(items,1)
    before=c.download('huaxing')[0];old=load_workbook(BytesIO(raw))
    for _ in range(3):
        h,snap,meta=latest();assert c.operate('huaxing','tester',replacement(before,version=h['revision']),[],'replace')['status']=='unchanged'
        assert run(items,h['revision'])['status']=='unchanged'
        got=c.download('huaxing')[0];assert got==before
    actual=book()
    for row in old['颜色和历史样例']:
        for cell in row:
            other=actual['颜色和历史样例'][cell.coordinate]
            assert other.value==cell.value and copy(other.fill)==copy(cell.fill) and copy(other.font)==copy(cell.font)
    _,snap,meta=latest();loc=meta['bindings'][items[0]['id']]
    assert actual[loc['sheet']].cell(loc['row'],10).fill.fgColor.rgb=='FF000000'
    assert actual[loc['sheet']].cell(loc['row'],10).font.color.rgb=='FFFFFFFF'
    assert len(snap)==1 and len(actual.sheetnames)==3


def test_unknown_black_is_preserved_and_warning_survives_reload():
    from openpyxl.styles import PatternFill,Font
    wb=load_workbook(BytesIO(colored_base()));wb['颜色和历史样例']['D5']='没有可追溯原件';wb['颜色和历史样例']['D5'].fill=PatternFill('solid',fgColor='FF000000');wb['颜色和历史样例']['D5'].font=Font(color='FF000000')
    stream=BytesIO();wb.save(stream);content=stream.getvalue()
    response=c.operate('huaxing','tester',replacement(content),[],'replace')
    assert '1 处' in response['warning'] and '颜色和历史样例!D5' in response['warning']
    assert c.download('huaxing')[0]==content
    assert c.state('huaxing',[],SCOPE)['warning']==response['warning']


def test_format_signature_mime_and_conversion_failure_preserve_head(monkeypatch):
    from app.services import qc_workbook_upload as upload
    rows=seed(1);run(rows);before=c.download('huaxing')[0];head=latest()[0]
    original=(c.TEMPLATES/'huaxing'/'original.xls').read_bytes()
    for name,content,mime in [('fake.xlsx',original,None),('fake.xls',colored_base(),None),('bad.xls',bytes.fromhex('D0CF11E0A1B11AE1')+b'broken',None),('bad.xlsx',b'broken',None),('valid.xlsx',colored_base(),'image/png')]:
        body=replacement(content,name,1)
        if mime:body['contentType']=mime
        with pytest.raises(HTTPException) as caught:c.operate('huaxing','tester',body,[],'replace')
        assert caught.value.status_code==422 and c.download('huaxing')[0]==before and latest()[0]==head
    monkeypatch.setenv('QC_XLS_CONVERTER','/missing/explicit/converter')
    with pytest.raises(HTTPException) as caught:c.operate('huaxing','tester',replacement(original,'年度.xls',1),[],'replace')
    assert caught.value.status_code==503 and c.download('huaxing')[0]==before


def test_native_xls_conversion_repairs_only_proven_fills_and_known_bad_reupload():
    from app.services import qc_workbook_upload as upload,qc_workbook_style as styles
    import xlrd
    from lxml import etree as E
    original=(c.TEMPLATES/'huaxing'/'original.xls').read_bytes();fixed,receipt=upload.prepare('原件.xls',original,c.template('huaxing'),'application/vnd.ms-excel')
    assert receipt['version']==2 and len(load_workbook(BytesIO(fixed)).sheetnames)==16
    # Reproduce the removed prior converter bug using the XLS XF context, without hard-coded cell addresses.
    p=writer.Package(fixed);sourcebook=xlrd.open_workbook(file_contents=original,formatting_info=True)
    fill=E.fromstring(b'<fill xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><patternFill patternType="solid"><fgColor rgb="FF000000"/></patternFill></fill>');fillid=styles.intern(p.styles,'fills',fill)
    for sh in sourcebook.sheets():
        nodes={cell.get('r'):cell for row in p.rows(sh.name) for cell in row}
        from openpyxl.utils import get_column_letter
        for r in range(sh.nrows):
            for col in range(sh.ncols):
                bg=sourcebook.xf_list[sh.cell_xf_index(r,col)].background;ref=get_column_letter(col+1)+str(r+1)
                if (bg.fill_pattern,bg.pattern_colour_index,bg.background_colour_index)!=(1,64,65) or ref not in nodes:continue
                cell=nodes[ref];xf=deepcopy(p.styles.find(writer.tag('cellXfs'))[int(cell.get('s','0'))]);xf.set('fillId',str(fillid));cell.set('s',str(styles.intern(p.styles,'cellXfs',xf)))
    polluted=p.finish();repaired,audit=styles.repair_fills(polluted,fixed,original)
    assert audit['count']==676
    before,after=load_workbook(BytesIO(polluted)),load_workbook(BytesIO(repaired))
    for name in before.sheetnames:
        assert before[name].merged_cells==after[name].merged_cells and before[name].sheet_state==after[name].sheet_state
        for row in before[name]:
            for cell in row:
                a,b=cell.value,after[name][cell.coordinate].value
                assert getattr(a,'text',a)==getattr(b,'text',b)
    c.operate('huaxing','tester',replacement(original,'原件.xls'),[],'replace')
    with c.store(True) as conn:
        h=conn.execute('SELECT * FROM heads').fetchone();v=conn.execute('SELECT * FROM versions WHERE id=?',(h['current_id'],)).fetchone();meta=json.loads(v['metadata']);meta['conversionReceipt']['version']=1
        conn.execute('UPDATE versions SET content=?,metadata=? WHERE id=?',(polluted,json.dumps(meta),v['id']))
    c.repair('huaxing','fixture-repair',1)
    result=c.operate('huaxing','tester',replacement(polluted,'已污染但可追溯.xlsx',2),[],'replace')
    assert result['status']=='replaced' and latest()[2]['styleRepair']['count']==676
    assert not styles.unresolved_black(c.download('huaxing')[0])


def test_owned_page_manual_content_blocks_replacement_atomically():
    rows=make_legacy_version()
    with c.store(True) as conn:
        h=conn.execute('SELECT * FROM heads').fetchone();v=conn.execute('SELECT * FROM versions WHERE id=?',(h['current_id'],)).fetchone();meta=json.loads(v['metadata']);p=writer.Package(v['content']);name=next(iter(meta['weekSheets'].values()));p.row(name,10000,['手工批注，不可丢弃'],{},20)
        conn.execute('UPDATE versions SET content=? WHERE id=?',(p.finish(),v['id']))
    before=c.download('huaxing')[0];head=latest()[0]
    with pytest.raises(HTTPException,match='独有内容'):c.repair('huaxing','tester',2)
    assert latest()[0]==head and c.download('huaxing')[0]==before


def test_new_year_known_xlsx_resets_old_bindings_without_source_deletion(monkeypatch):
    from datetime import datetime
    rows=seed(1);c.operate('huaxing','tester',replacement(colored_base()),[],'replace');run(rows,1);old=c.download('huaxing')[0]
    class NextYear(datetime):
        @classmethod
        def now(cls,tz=None):return cls(2027,1,1,tzinfo=tz)
    monkeypatch.setattr(c,'datetime',NextYear)
    response=c.operate('huaxing','tester',replacement(old,version=2),[],'replace')
    assert response['status']=='replaced' and not response['undoAvailable']
    h,snap,meta=latest();assert snap=={} and meta['annualYear']==2027
    assert sum(cell.value==rows[0]['poNo'] for row in book()['ZURU行验货'] for cell in row)==1
    assert book()['验货明细'].max_row==1
    with pytest.raises(HTTPException,match='跨年度'):run(rows,3)
    with sheet.store() as conn:assert sheet.get(conn,'huaxing',rows[0]['id']) is not None
    rows[0]['dateFrom']=rows[0]['dateTo']='2027-01-04';assert run(rows,3,scope={'dateFrom':'2027-01-01','dateTo':'2027-12-31'})['status']=='imported'


def test_more_than_one_page_complete_saved_snapshot_and_readonly_download():
    items=seed(80)
    for i,item in enumerate(items):items[i]=save(item,['PASS','FAIL','PENDING',''][i%4],('保留中文换行及完整前导零订单\\n'*10)+str(i))
    run(items);_,snap,meta=latest();w=book();assert len(snap)==80 and w['验货明细'].max_row==81
    assert set(meta['detailIds'])=={item['id'] for item in items}==set(meta['bindings'])
    for item in items:
        binding=meta['bindings'][item['id']];summary=meta['summaryRows'][item['id']];second=binding['sheet'] in writer.SECOND
        assert (w[binding['sheet']].cell(binding['row'],8 if second else 9).value or '')==item['result']
        assert (w['验货明细'].cell(summary['row'],14).value or '')==item['result']
        assert w['验货明细'].cell(summary['row'],15).value==item['remark']
    before=c.download('huaxing')[0];items[0]=save(items[0],'FAIL','此变更尚未重导')
    assert c.download('huaxing')[0]==before and latest()[1][items[0]['id']]['remark']!='此变更尚未重导'


def test_added_customer_second_layout_retains_its_own_columns_on_reimport():
    wb=load_workbook(BytesIO(colored_base()));ws=wb['ZURU行验货'];ws.title='新客户'
    for i,value in enumerate(writer.HEAD2,1):ws.cell(1,i,value)
    stream=BytesIO();wb.save(stream);c.operate('huaxing','tester',replacement(stream.getvalue()),[],'replace')
    src=source('custom-layout','000 NEW-123');src.update(customer='新客户',country='USA',dateFrom='2026-10-29',dateTo='2026-10-29')
    rows=sheet.sync('huaxing',pushed_candidates([src]),False)['items'];run(rows,1)
    rows[0]=save(rows[0],'FAIL','第二次导入保留原布局');run(rows,2)
    _,snap,meta=latest();loc=meta['bindings'][rows[0]['id']];got=book()['新客户'];assert '新客户' not in meta.get('createdCustomers',[])
    assert got.cell(loc['row'],1).value=='000 NEW-123' and got.cell(loc['row'],2).value=='USA'
    assert got.cell(loc['row'],8).value=='FAIL' and got.cell(loc['row'],11).value=='第二次导入保留原布局'


def test_source_query_order_only_does_not_create_version_or_replace_undo():
    rows=seed(3);run(rows);before=c.download('huaxing')[0];head=latest()[0]
    assert run(list(reversed(rows)),1)['status']=='unchanged'
    assert c.download('huaxing')[0]==before and latest()[0]==head


def test_new_date_headers_have_space_without_changing_historical_columns():
    from datetime import datetime
    rows=seed(3);before=load_workbook(BytesIO(c.download('huaxing')[0]));run(rows)
    _,_,meta=latest();after=load_workbook(BytesIO(c.download('huaxing')[0]))
    for name,owned in meta['ownedRows'].items():
        dates=[ref for ref in meta['ownedMerges'].get(name,[]) if ref.startswith('A') and ':C' in ref]
        if owned:assert dates
        for ref in dates:assert after[name][ref.split(':')[0]].value is not None
        dims=lambda ws:[(k,d.width,d.min,d.max,d.hidden,d.bestFit,d.outlineLevel,d.collapsed) for k,d in ws.column_dimensions.items()]
        if name in before.sheetnames:assert dims(before[name])==dims(after[name])
    c.operate('huaxing','test',{'requestKey':uuid4().hex,'expectedVersion':1},[],'undo')
    restored=load_workbook(BytesIO(c.download('huaxing')[0]))
    for name in before.sheetnames:assert before[name].merged_cells==restored[name].merged_cells
