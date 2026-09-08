import re
import json
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from openpyxl.utils.cell import range_boundaries

from app.services.internal_quote import _rr2_cost_summary
from app.services.internal_quote_calculator import calculate_section
from app.services.internal_quote_excel import _build_summary_sheet, _purchase_source_lines, _pricing_entry_amount
from test_internal_quote_excel_formula_links import _section, _find_row


def arithmetic_value(sheet, coordinate):
    """Evaluate the small arithmetic/SUM subset used by visible cost blocks."""
    value = sheet[coordinate.replace('$', '')].value
    if not isinstance(value, str) or not value.startswith('='):
        return float(value or 0)
    expression = value[1:].replace('$', '')

    def sum_range(match):
        left, top, right, bottom = range_boundaries(match[1])
        return str(sum(arithmetic_value(sheet, sheet.cell(row, col).coordinate)
                       for row in range(top, bottom + 1) for col in range(left, right + 1)))

    expression = re.sub(r'SUM\(([A-Z]+\d+:[A-Z]+\d+)\)', sum_range, expression)
    expression = re.sub(r'\b[A-Z]+\d+\b', lambda match: str(arithmetic_value(sheet, match[0])), expression)
    assert re.fullmatch(r'[0-9.eE+*/()\s-]+', expression), expression
    return eval(expression, {'__builtins__': {}}, {})


@pytest.mark.parametrize('component', [False, True])
def test_original_purchase_prices_quantities_losses_and_detached_groups(component):
    snapshot = {'fx': {'rmb_hkd': '.85', 'hkd_usd': '7.8', 'rmb_usd': '6.63'}, 'electronic_profit_rate_percent': '0'}
    owner = {'pricing_component_id': 'body'} if component else {}
    override = {} if component else {'markup_override': '1.08'}
    materials = [
        {'item': '弹簧', 'category': 'hardware', 'quantity': 2, 'unit_price_rmb': '.52', **owner},
        {'item': '螺丝', 'category': 'hardware', 'quantity': 1, 'unit_price_rmb': '.13', **owner},
        {'item': '镜片', 'category': 'auxiliary', 'auxiliary_category': '其他外购', 'quantity': 2,
         'unit_price_rmb': '999', 'unit_price_hkd': '.8', 'unit_price_source_currency': 'HKD', 'loss_rate': '1.02', **owner, **override},
        {'item': '内垫', 'category': 'auxiliary', 'quantity': 1, 'unit_price_rmb': '.123456', 'loss_rate': '1.03', **owner},
    ]
    sales = {'packaging_materials': [
        {'item': '胶袋', 'category': 'other_purchase', 'quantity': 3, 'unit_price_rmb': '.17', 'loss_rate': '1.02'},
        {'item': '说明书', 'category': 'leaflet_manual', 'quantity': 1, 'unit_price_hkd': '.2', 'unit_price_source_currency': 'HKD'},
    ], 'shipping': {'markup_x': '1.2', 'misc_ratio': '.02'}, 'testing_fee_enabled': False, 'freight_calc': {'enabled': False},
        'cartons': [{'item': '外箱', 'length_in': 8, 'width_in': 6, 'height_in': 5, 'qty_per_carton': 12}], 'paper_price_factor': '2.75',
        'color_box_size_in': {'length': 7.25, 'width': 5.25, 'height': 4}}
    if component:
        sales.update(pricing_mode='component', pricing_components=[{'id': 'body', 'name': '主体', 'markup_x': '1.15'}])
    sections = []
    for code, payload in [('engineering', {'materials': materials}), ('sales', sales)]:
        calc = calculate_section(code, payload, snapshot, 'PURCHASE-EXPORT')
        assert calc['status'] == 'valid', calc.get('warnings')
        section = _section(code, payload=payload, calculation=calc)
        section.is_required = True
        sections.append(section)
    totals = [json.loads(section.calculation_json)['totals'] for section in sections]
    context = {key: Decimal(totals[0][key]) for key in ('hardware_hkd', 'auxiliary_hkd')}
    context.update({key: Decimal(totals[1][key]) for key in ('packaging_material_hkd', 'carton_hkd')})
    context['factory_price_hkd'] = sum(context.values())
    summary = _rr2_cost_summary(sections, context, snapshot, factory_id='huakang-b' if component else 'huaxing')
    workbook = Workbook()
    _build_summary_sheet(workbook, SimpleNamespace(product_name='采购换算验证', quote_no='PURCHASE', customer='JustPlay' if component else '普通客', region_code='mainland', remark=''),
                         sections, snapshot, summary, context, [])
    sheet = workbook['报价明细']
    expected = {'弹簧': '=0.52/$L$4*2', '螺丝': '=0.13/$L$4', '镜片': '=0.8*2*1.02',
                '内垫': '=0.123456/$L$4*1.03', '胶袋': '=0.17/$L$4*3*1.02', '说明书': .2}
    for name, formula in expected.items():
        values = [sheet.cell(row, 4).value for row in range(1, sheet.max_row + 1) if sheet.cell(row, 3).value == name]
        assert formula in values or (component and isinstance(formula, str) and any(
            isinstance(value, str) and value.startswith(formula + '*') for value in values)), (name, values)
    if not component:
        # The detached mirror cost is removed once from the main cost range.
        mirror_rows = [row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 3).value == '镜片']
        assert len(mirror_rows) == 2
        assert arithmetic_value(sheet, f'D{mirror_rows[0]}') == 0
        assert arithmetic_value(sheet, f'D{mirror_rows[1]}') == pytest.approx(1.632)
    workbook.close()


@pytest.mark.parametrize('mode', ['detail', 'quick', 'legacy'])
def test_electronic_source_currency_is_based_on_entered_price_and_quick_qty_is_one(mode):
    rows = [{'item': 'IC', 'quantity': 2, 'unit_price_rmb': '.52'}, {'item': 'LED', 'quantity': 3, 'unit_price_hkd': '.2'}]
    payload = {'components': rows}
    if mode == 'quick':
        payload = {'quote_mode': 'quick', 'quick_quotes': rows}
    elif mode == 'legacy':
        rows[0] = {'item': 'IC', 'quantity': 2, 'unit_price_hkd': '.52'}
    calc = calculate_section('electronic', payload, {'fx': {'rmb_hkd': '.85'}}, 'SOURCE')
    lines = _purchase_source_lines(_section('electronic', payload=payload, calculation=calc))
    assert [_pricing_entry_amount(line) for line in lines] == (
        ['=0.52/$L$4', .2] if mode == 'quick' else
        ['=0.52*2', '=0.2*3'] if mode == 'legacy' else ['=0.52/$L$4*2', '=0.2*3'])


def test_detached_electronic_purchase_moves_allocated_overhead_without_negative_main_price():
    snapshot = {'fx': {'rmb_hkd': '.85', 'hkd_usd': '7.8'}}
    payload = {'components': [{'item': 'IC', 'quantity': 2, 'unit_price_rmb': '.52', 'markup_override': '1.08'},
                              {'item': 'LED', 'quantity': 3, 'unit_price_hkd': '.2'}], 'profit_rate_percent': 10}
    calc = calculate_section('electronic', payload, snapshot, 'DETACHED')
    sections = [_section('electronic', payload=payload, calculation=calc),
                _section('sales', payload={'shipping': {'markup_x': '1.2'}})]
    for section in sections:
        section.is_required = True
    total = Decimal(calc['totals']['total_hkd'])
    context = {'electronic_hkd': total, 'factory_price_hkd': total}
    summary = _rr2_cost_summary(sections, context, snapshot)
    workbook = Workbook()
    _build_summary_sheet(workbook, SimpleNamespace(product_name='电子倍率验证', quote_no='DETACHED', customer='普通客', region_code='mainland', remark=''),
                         sections, snapshot, summary, context, [])
    sheet = workbook['报价明细']
    rows = [row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 3).value == 'IC']
    assert len(rows) == 2
    assert arithmetic_value(sheet, f'D{rows[0]}') == 0
    assert sheet.cell(rows[1], 4).value.startswith('=0.52/$L$4*2*')
    assert arithmetic_value(sheet, f'D{rows[1]}') + arithmetic_value(sheet, f'D{_find_row(sheet, 3, "LED")}') == pytest.approx(float(total), abs=.0005)
    workbook.close()
