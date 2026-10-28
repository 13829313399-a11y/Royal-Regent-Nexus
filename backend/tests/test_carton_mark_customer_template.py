import base64
import io
import json

import httpx
import pytest
from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image
from pypdf import PdfReader
from pydantic import SecretStr

from app.services import carton_mark_customer_template as template
from app.services import carton_mark_flexible as flow
from app.services import carton_mark_images as images
from app.services import carton_mark_workbench as workbench
from app.services import carton_mark_ai_pilot as pilot
from app.services import carton_supplier_mark_ai as ai


def original(po='TEST-PO-NEW', with_image=False):
    book=Workbook();sheet=book.active;sheet.title='本单'
    values={'A1':'正唛 X2','H1':'侧唛 X2','M2':'备注：本单侧唛文字居中；两面ZURU LOGO和向上标识。',
        'A4':'SUPPLIER CODE:', 'B4':'TEST-114', 'A5':'SMYTHS ORDER NO:', 'B5':po,
        'A6':'ADDRESS:', 'B6':'TEST CUSTOMER ADDRESS', 'A7':'SUPPLIER ITEM NO:', 'B7':'TEST-ITEM-NEW',
        'A8':"SMYTHS ITEM NO:", 'B8':'TEST-ITEM-02', 'A9':'DESCRIPTION:', 'B9':'NEW PRODUCT',
        'A10':'QUANTITY:', 'B10':'4 PCS', 'A11':'MADE IN:', 'B11':'CHINA',
        'A12':'LANGUAGES:', 'B12':'EN, FR, DE, NL, IT', 'A13':'OUTER CARTON BARCODE:', 'B13':'10193052021172',
        'A15':'TEST FRANCE ADDRESS care@example.invalid',
        'H4':'G.W.:','I4':'(按实际填写) KGS','H5':'N.W.:','I5':'(按实际填写) KGS','H6':'MEAS.:','I6':'X X CM'}
    for key,value in values.items():sheet[key]=value
    if with_image:
        data=io.BytesIO();Image.new('RGB',(320,180),'white').save(data,format='PNG');data.seek(0)
        sheet.add_image(ExcelImage(data),'M5')
    out=io.BytesIO();book.save(out);book.close()
    return dict(excel_file_name='new-order.xlsx',excel_bytes=out.getvalue(),customer_template_id=template.ID,image_consent=True)


def response(source):
    cells={c.location.split('!')[-1]:c for c in source['cells']}
    def selection(loc):return dict(source_id=cells[loc].id,value=cells[loc].text)
    def field(a,b=None,role='normal',position='before_barcode'):
        return dict(label=selection(a) if b else None,parts=[selection(b or a)],role=role,position=position)
    front=[field('A'+str(r),'B'+str(r)) for r in range(4,13)]+[field('A15',position='after_barcode')]
    side=[field('H'+str(r),'I'+str(r),role='placeholder') for r in range(4,7)]
    note=next(n.id for n in source['notes'] if n.kind=='current_note')
    return dict(template_id='source-flow-v1',paper='A4',orientation='landscape',records=[dict(panels=[
        dict(face='front',copies=2,logo=True,fields=front,barcodes=[dict(source=selection('B13'),kind='ITF14',width_mm=50.,height_mm=15.,frame='box',note_ids=[])],symbols=[dict(kind='this_side_up',note_ids=[note])]),
        dict(face='side',copies=2,logo=True,fields=side,barcodes=[],symbols=[dict(kind='this_side_up',note_ids=[note])])])],
        notes=[],ignored_cells=[],template_changes=[dict(record=0,face='side',note_ids=[note],summary='本单侧唛文字居中',align='center')])


def setup(monkeypatch, source=None):
    source=source or original()
    monkeypatch.setattr(ai,'provider_configuration',lambda:({},'test-provider'))
    monkeypatch.setattr(images,'_ocr',lambda serialized:[])
    calls=[]
    def provider(evidence,*args,**kwargs):
        calls.append((evidence,kwargs))
        return json.dumps(response(flow.read(source['excel_file_name'],source['excel_bytes'])),ensure_ascii=False),{}
    monkeypatch.setattr(pilot,'request_plan',provider)
    return source,calls


def render(source,result,**kwargs):
    return workbench.render(source,'owner',result['payload'],result['signature'],kwargs.pop('edits',{}),[],
                            kwargs.pop('paper','A4'),'landscape',include_sources=kwargs.pop('include_sources',False),**kwargs)


def test_customer_template_sends_images_and_notes_and_never_sends_old_business_values(monkeypatch):
    source,calls=setup(monkeypatch,original(with_image=True))
    before=template.PROFILE.read_bytes()
    result=workbench.read(source,'owner',True)
    assert result['mode']!='needs_mapping', result['warnings']
    assert result['customer_template']['id']==template.ID
    assert result['order_changes'][0]['align']=='center'
    evidence,request=calls[0]
    assert {i['id'] for i in request['images']}=={'I1','T1'}
    assert any(e.kind=='current_note' for e in evidence)
    assert all('_pixels' not in i for i in result['images'])
    assert request['images'][0]['data'] not in result['payload']
    pdf,pages,_=render(source,result)
    text=PdfReader(io.BytesIO(pdf)).pages[0].extract_text()
    assert pages==1 and 'TEST-PO-NEW' in text and 'NEW PRODUCT' in text
    assert '4500227434' not in text and '4400610360' not in text and '77903TQ1' not in text
    render(source,result,paper='A3')
    assert len(calls)==1 and template.PROFILE.read_bytes()==before
    full,_,_=render(source,result,include_sources=True)
    assert '本单侧唛文字居中' in '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(full)).pages)


def test_customer_template_requires_image_consent_and_valid_profile_before_provider(monkeypatch):
    source,calls=setup(monkeypatch)
    for update in [{'image_consent':False},{'customer_template_id':'unknown'}]:
        with pytest.raises(HTTPException):workbench.read({**source,**update},'owner',True)
    with pytest.raises(HTTPException):workbench.read(source,'owner',False)
    assert not calls


def test_explicit_template_artwork_can_be_adjusted_without_brand_text_in_original(monkeypatch):
    from openpyxl import load_workbook
    source=original();book=load_workbook(io.BytesIO(source['excel_bytes']))
    book.active['M2']='备注：本单侧唛文字居中。'
    data=io.BytesIO();book.save(data);book.close();source['excel_bytes']=data.getvalue()
    setup(monkeypatch,source)
    plan=response(flow.read(source['excel_file_name'],source['excel_bytes']))
    for panel in plan['records'][0]['panels']:
        panel['logo']=False;panel['symbols']=[]
    monkeypatch.setattr(pilot,'request_plan',lambda *a,**kw:(json.dumps(plan,ensure_ascii=False),{}))
    result=workbench.read(source,'owner',True)
    assert result['mode']=='ai' and result['logo_available']
    pdf,_,_=render(source,result,graphics=[dict(record=0,face='side',logo=True,symbols=['keep_dry'])])
    assert PdfReader(io.BytesIO(pdf)).pages[0].images


def test_order_exception_cannot_transfer_or_mutate_template_and_revision_invalidates_read(monkeypatch):
    source,calls=setup(monkeypatch);result=workbench.read(source,'owner',True)
    with pytest.raises(HTTPException):render(original('ANOTHER-ORDER'),result)
    with pytest.raises(HTTPException):
        workbench.render(source,'other-owner',result['payload'],result['signature'],{},[],'A4','landscape')
    monkeypatch.setattr(template,'signature',lambda:'new-version')
    with pytest.raises(HTTPException):render(source,result)
    assert len(calls)==1


def test_order_exception_can_be_declined_locally_and_identity_edits_stay_locked(monkeypatch):
    source,calls=setup(monkeypatch);result=workbench.read(source,'owner',True)
    locked=[f for p in result['records'][0]['panels'] for f in p['fields'] if f['locked']]
    assert any(f['label']=='SMYTHS ORDER NO:' for f in locked)
    for field in locked:
        with pytest.raises(HTTPException):render(source,result,edits={field['id']:'FORGED'})
        with pytest.raises(HTTPException):render(source,result,omissions=[field['id']])
    pdf,_,_=render(source,result,include_sources=True,ignored_changes=[0])
    assert '本次未采用此修改' in '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    with pytest.raises(HTTPException):render(source,result,ignored_changes=[True])
    with pytest.raises(HTTPException):render(source,result,ignored_changes=[99])
    assert len(calls)==1


@pytest.mark.parametrize('patch',[
    {'note_ids':['E1']},{'note_ids':['T1']},{'record':99},{'font_pt':99.0},
    {'quantity':'999 PCS'},{'logo_visible':'false'},
])
def test_only_bounded_source_cited_layout_exceptions_are_allowed(patch):
    source=flow.read('new.xlsx',original()['excel_bytes']);data=response(source)
    plan=flow.Plan.model_validate({k:v for k,v in data.items() if k!='template_changes'})
    data['template_changes'][0].update(patch)
    with pytest.raises(pilot.PilotError):template.changes(json.dumps(data),plan,source)


def test_image_transport_contains_real_pixels_and_rejects_urls_or_oversize_without_request():
    picture=io.BytesIO();Image.new('RGB',(20,20),'white').save(picture,format='JPEG')
    entry=dict(id='I1',location='本单!image:1',data=base64.b64encode(picture.getvalue()).decode())
    evidence=[pilot.Evidence(id='I1',kind='image_reference_text',location=entry['location'],text='备注图片')]
    connection=dict(base_url='https://api.deepseek.com',protocol='openai',model='deepseek-flash',api_key=SecretStr('test-only'))
    calls=[]
    def handler(request):
        calls.append(json.loads(request.content));return httpx.Response(200,json={'choices':[{'finish_reason':'stop','message':{'content':'{}'}}]})
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        pilot.request_plan(evidence,flow.TEMPLATE,connection,client=client,images=[entry])
        for bad in [[{**entry,'data':'https://example.invalid/private'}],[entry]*13,[{**entry,'id':'unknown'}]]:
            with pytest.raises(pilot.PilotError):pilot.request_plan(evidence,flow.TEMPLATE,connection,client=client,images=bad)
    parts=calls[0]['messages'][1]['content']
    assert len(calls)==1 and parts[-1]['type']=='image_url' and entry['data'] in parts[-1]['image_url']['url']


def test_existing_text_only_consent_does_not_enable_pixels(monkeypatch):
    monkeypatch.setattr(images,'_ocr',lambda serialized:[])
    source=original(with_image=True)
    _,public,_=images.collect(source['excel_file_name'],source['excel_bytes'])
    assert all('_pixels' not in i for i in public)


def test_multipart_read_and_render_transport_with_isolated_router(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from types import SimpleNamespace
    from app.api import carton_supplier_portal as api
    source,calls=setup(monkeypatch)
    user=SimpleNamespace(id='test-user',authorization_version=1)
    application=FastAPI();application.include_router(api.router)
    application.dependency_overrides[api.get_current_user]=lambda:user
    application.dependency_overrides[api.get_db]=lambda:None
    async def input_source(*args):
        return {},None,{k:v for k,v in source.items() if k not in {'image_consent','customer_template_id'}},'owner'
    monkeypatch.setattr(api,'_carton_review_input',input_source)
    monkeypatch.setattr(api,'_carton_review_recheck',lambda *a:user)
    with TestClient(application) as client:
        base='/api/carton-supplier/carton-mark/'
        data=dict(ai_consent='true',image_consent='true',customer_template_id=template.ID)
        read=client.post(base+'read-pdf-draft',data=data)
        assert read.status_code==200,read.text
        value=read.json();assert value['customer_template']['id']==template.ID
        rendered=client.post(base+'render-reviewed-draft',data=dict(review_payload=value['payload'],review_signature=value['signature'],paper='A4',orientation='landscape',include_sources='false'))
        assert rendered.status_code==200,rendered.text
        assert rendered.headers['cache-control']=='private, no-store' and len(calls)==1
        assert 'TEST-PO-NEW' in PdfReader(io.BytesIO(rendered.content)).pages[0].extract_text()
