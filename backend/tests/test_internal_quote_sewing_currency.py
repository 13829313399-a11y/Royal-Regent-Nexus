import json
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.services.internal_quote_calculator import CalculationInputError, calculate_section
from app.services.internal_quote_excel import _build_sewing_sheet
from app.services.internal_quote_import import parse_internal_quote_workbook


SNAPSHOT = {"fx": {"rmb_hkd": "0.85"}}


def parse_rows(header, rows):
    book = Workbook()
    book.active.append(header)
    for row in rows:
        book.active.append(row)
    output = BytesIO()
    book.save(output)
    return parse_internal_quote_workbook(output.getvalue(), "sewing").payload_fragment


def calculate(payload):
    return calculate_section("sewing", payload, SNAPSHOT, "test")


def test_legacy_hkd_source_reconciles_material_and_labor_without_second_conversion():
    # Representative cached source values include prices already converted to HKD.
    payload = parse_rows(
        ["布料名称", "裁片部位", "用量/码", "单价", "码点", "价钱"],
        [["布标", "", 1, 0.235294117647059, 1.1, 0.258823529411765],
         ["裁床人工", "", 1, 0.56470588235294, 1.1, 0.621176470588234],
         ["车缝人工", "", 1, 4.05882352941176, 1.1, 4.46470588235294]],
    )
    assert payload["groups"][0]["materials"][0]["unit_price_hkd"] == "0.235294117647059"
    result = calculate(payload)
    assert result["totals"]["total_hkd"] == "5.3447"
    assert result["totals"]["material_hkd"] == "0.2588"
    assert result["totals"]["labor_hkd"] == "5.0859"
    assert result["totals"]["total_rmb"] == "4.5430"


@pytest.mark.parametrize("unit_header,rate_column,currency,expected", [
    ("单价 RMB", False, "RMB", "2.0000"),
    ("单价", True, "RMB", "2.0000"),
    ("单价 HKD", True, "HKD", "1.7000"),
    ("单价(港币)", False, "HKD", "1.7000"),
    ("单价", False, "HKD", "1.7000"),
])
def test_import_respects_unit_currency_and_rate_template(unit_header, rate_column, currency, expected):
    header = ["物料名称", "部位", "用量", unit_header, "码点", "价钱"]
    row = ["布", "身体", 1, 1.7, 1, 2]
    if rate_column:
        header.append("汇率")
        row.append(0.85)
    payload = parse_rows(header, [row])
    assert payload["groups"][0]["materials"][0]["unit_price_source_currency"] == currency
    assert calculate(payload)["totals"]["total_hkd"] == expected


def test_missing_unit_uses_total_currency_and_does_not_apply_markup_twice():
    payload = parse_rows(
        ["物料名称", "部位", "用量", "单价(HKD)", "码点", "价钱(HKD)"],
        [["人工", "", 2, None, 1.1, 4.4]],
    )
    assert payload["groups"][0]["materials"][0]["unit_price_hkd"] == "2.0000"
    assert calculate(payload)["totals"]["total_hkd"] == "4.4000"


def test_explicit_rmb_total_identifies_no_rate_sheet_but_hkd_total_does_not_override_rate_template():
    payload = parse_rows(
        ["物料名称", "部位", "用量", "单价", "码点", "价钱(RMB)"],
        [["布", "", 1, 1.7, 1, 1.7]],
    )
    assert calculate(payload)["totals"]["total_hkd"] == "2.0000"
    payload = parse_rows(
        ["物料名称", "部位", "用量", "单价", "码点", "价钱 HKD", "汇率"],
        [["布", "", 1, 1.7, 1, 2, .85]],
    )
    assert calculate(payload)["totals"]["total_hkd"] == "2.0000"


def test_hkd_source_wins_over_stale_rmb_and_legacy_rmb_keeps_existing_behavior():
    payload = {"groups": [{"materials": [
        {"item": "布", "usage": 2, "unit_price_hkd": "3", "unit_price_rmb": "99", "unit_price_source_currency": "HKD", "exchange_rate": ".8", "markup": "1.1"},
        {"item": "人工", "usage": 1, "unit_price_rmb": "1.7", "markup": 1},
    ]}]}
    result = calculate(payload)
    assert result["totals"]["total_hkd"] == "8.6000"
    assert result["totals"]["total_rmb"] == "6.9800"
    assert result["line_breakdown"][0]["formula"] == "usage * unit_price_hkd * markup"
    payload["groups"][0]["materials"][0]["unit_price_source_currency"] = "USD"
    with pytest.raises(CalculationInputError, match="RMB 或 HKD"):
        calculate(payload)


@pytest.mark.parametrize("currency", ["HKD", " hkd "])
def test_export_retains_original_prices_currency_and_mixed_currency_reimport(currency):
    payload = {"groups": [{"name": "衣服", "materials": [
        {"item": "布", "usage": 2, "unit_price_hkd": "3.123456789", "unit_price_source_currency": currency, "markup": 1.1},
        {"item": "人工", "usage": 1, "unit_price_rmb": "1.7", "markup": 1},
    ]}]}
    book = Workbook()
    _build_sewing_sheet(book, SimpleNamespace(payload_json=json.dumps(payload)), SNAPSHOT)
    output = BytesIO()
    book.save(output)
    exported = load_workbook(BytesIO(output.getvalue()))["车缝明细"]
    assert exported["G5"].value == 3.123456789
    assert exported["M5"].value == "HKD"
    assert exported["M6"].value == "RMB"
    assert exported["I5"].value == '=IF(M5="HKD",F5*G5,F5*G5/H5)'
    imported = parse_internal_quote_workbook(output.getvalue(), "sewing").payload_fragment
    assert calculate(imported)["totals"] == calculate(payload)["totals"]


def test_hkd_import_preview_confirm_and_reload_retain_currency_and_amount(monkeypatch):
    from test_internal_quote_api import ALL_SECTION_CODES, create_payload, login, make_client

    with make_client(monkeypatch) as client:
        login(client, "sewing-currency-owner", "admin", "*", "*")
        created = client.post("/api/internal-quotes", json=create_payload(
            suffix="SEWING-HKD", participating_sections=ALL_SECTION_CODES,
        ))
        assert created.status_code == 201, created.text
        quote_id = created.json()["id"]
        book = Workbook()
        book.active.append(["物料名称", "部位", "用量", "单价", "码点", "价钱"])
        book.active.append(["布标", "", 1, .235294117647059, 1.1, .258823529411765])
        output = BytesIO()
        book.save(output)
        preview = client.post(f"/api/internal-quotes/{quote_id}/imports/sewing/preview", files={
            "file": ("sewing.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        })
        assert preview.status_code == 201, preview.text
        batch = preview.json()
        response = client.post(f"/api/internal-quotes/{quote_id}/imports/{batch['batch_id']}/confirm", json={
            "revision": batch["target_revision"], "mode": "replace",
        })
        assert response.status_code == 200, response.text
        sections = client.get(f"/api/internal-quotes/{quote_id}").json()["sections"]
        sewing = next(row for row in sections if row["department"] == "sewing")
        material = sewing["payload"]["groups"][0]["materials"][0]
        assert material["unit_price_source_currency"] == "HKD"
        assert material["unit_price_hkd"] == "0.235294117647059"
        assert sewing["calculation"]["totals"]["total_hkd"] == "0.2588"
