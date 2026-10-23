from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook

from test_molding_sample_api import make_client, login_as
from test_carton_transaction_guards_api import BASE, ledger
from test_carton_supplier_settlement_api import clock, view, payload as bill_payload, action, close_month


def unordered(unit="cm", **line_changes):
    return {"factory_id": "huaxing", "post_immediately": True, "request_id": uuid4().hex,
            "delivery_note_no": uuid4().hex, "delivery_date": "2026-10-02", "acceptance_date": "2026-10-02",
            "lines": [{"source_type": "AD_HOC", "unordered_receipt": True, "dimension_unit": unit,
                       "packaging_type": "普通箱", "paper_quality": "B3B", "specification": "10 × 20 × 30",
                       "delivered_quantity": "10", "received_quantity": "10", "unit_price": "2", "location": "TEST/A", **line_changes}]}


def test_new_unordered_receipts_cannot_create_inventory_or_monthly_payables(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "carton_supervisor"); clock(monkeypatch)
        place = client.post(BASE + "/inventory/locations", json={"factory_id": "huaxing", "warehouse": "TEST", "bin_code": "A"})
        assert place.status_code == 201, place.text
        for unit in ("cm", "in"):
            body = unordered(unit, location_allocations=[{"location_id": place.json()["id"], "quantity": "10"}])
            response = client.post(BASE + "/receipts", json=body)
            assert response.status_code == 422 and "来源" in response.text
        assert ledger(client) == []
        assert view(client)["sources"] == []
        assert client.get(BASE + "/orders", params={"factory_id": "huaxing"}).json()["total"] == 0


def test_unordered_explicit_opt_in_required_fields_and_scope(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "warehouse_keeper"); clock(monkeypatch)
        for changes in [{"unordered_receipt": False}, {"source_type": "FORMAL_ORDER"}, {"unit_price": "0"}, {"dimension_unit": "mm"}, {"packaging_type": ""}, {"specification": "10 in"}]:
            response = client.post(BASE + "/receipts", json=unordered(**changes))
            assert response.status_code == 422, response.text
        missing_date = unordered(); missing_date.pop("acceptance_date")
        assert client.post(BASE + "/receipts", json=missing_date).status_code == 422
        wrong_factory = {**unordered(), "factory_id": "huadeng"}
        assert client.post(BASE + "/receipts", json=wrong_factory).status_code == 403
        login_as(client, "engineer")
        assert client.post(BASE + "/receipts", json=unordered()).status_code == 403
        login_as(client, "warehouse_keeper")
        assert ledger(client) == []
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonSupplier
        with SessionLocal() as db:
            supplier = CartonSupplier(id="OTHER-SUPPLIER", factory_id="huaxing", supplier_code="OTHER",
                supplier_name="Other", status="ACTIVE", created_at="2026-10-01", updated_at="2026-10-01")
            db.add(supplier); db.commit()
        response = client.post(BASE + "/receipts", json={**unordered(), "supplier_id": "OTHER-SUPPLIER"})
        assert response.status_code == 422 and "来源" in response.text
        from test_carton_procurement_api import _ensure_dickie_customer
        _ensure_dickie_customer(client)
        regular = unordered(unordered_receipt=False, customer_code="DICKIE", item_no="TEST-REGULAR", location="A-01")
        response = client.post(BASE + "/receipts", json=regular)
        assert response.status_code == 422 and "来源" in response.text
        assert ledger(client) == []


def test_delivery_excel_material_only_rows_and_pdf_ocr_without_identifiers(monkeypatch):
    from app.services import carton_procurement_imports as imports
    workbook = Workbook(); sheet = workbook.active
    sheet.append(["纸质", "规格", "数量"])
    sheet.append(["B3B 普通箱", "10*20*30", 12])
    data = BytesIO(); workbook.save(data)
    result = imports._parse_delivery_spreadsheet("material.xlsx", data.getvalue())
    assert len(result["rows"]) == 1
    assert result["rows"][0]["item_no"] == result["rows"][0]["contract_no"] == ""
    assert result["rows"][0]["delivered_quantity"] == 12
    assert result["rows"][0]["specification"] == "10*20*30"
    sheet.delete_cols(2); sheet.insert_cols(2, 3)
    sheet.cell(1, 2, "长"); sheet.cell(1, 3, "宽"); sheet.cell(1, 4, "高")
    sheet.cell(2, 2, 10); sheet.cell(2, 3, 20); sheet.cell(2, 4, 30)
    data = BytesIO(); workbook.save(data)
    result = imports._parse_delivery_spreadsheet("dimensions.xlsx", data.getvalue())
    assert not result["rows"][0]["specification"].endswith("in")
    monkeypatch.setattr(imports, "_document_text", lambda *args: ("普通箱 B3B 10x20x30 cm 12 PCS", "test", ""))
    result = imports._parse_delivery_document("material.pdf", b"test")
    assert len(result["rows"]) == 1 and result["rows"][0]["delivered_quantity"] == 12
    assert result["rows"][0]["contract_no"] == result["rows"][0]["item_no"] == ""


@pytest.mark.parametrize("spec,field,expected", [
    ("12 × 10 × 5", "", "cm"), ("12 × 10 × 5 in", "", "in"),
    ("12 × 10 × 5 inch", "", "in"), ("12 × 10 × 5 [in]", "", "in"),
    ("12 × 10 × 5 英寸", "", "in"), ("12 × 10 × 5 in", "cm", "cm"),
])
def test_legacy_display_dimension_without_rewriting_identity(spec, field, expected):
    from app.services.carton_dimensions import dimension_unit_for_display, specification_for_display
    assert dimension_unit_for_display(field, spec) == expected
    display_unit = "inch" if expected == "in" else expected
    assert specification_for_display(spec, field) == f"12 × 10 × 5 {display_unit}"


def test_unordered_unit_normalization_rejects_unsupported_or_conflicting_units():
    from app.services.carton_procurement import _unordered_specification
    from fastapi import HTTPException
    assert _unordered_specification("12*10*5 英寸", "in") == "12 × 10 × 5 [in]"
    assert _unordered_specification("12*10*5 [cm]", "cm") == "12 × 10 × 5 [cm]"
    assert _unordered_specification("12*10*5cm", "cm") == "12 × 10 × 5 [cm]"
    for spec in ["12*10*5 mm", "12*10*5 毫米", "12*10*5 [mm]", "12*10*5 ft", "12*10*5 foo", "12*10*5 mm [cm]", "12*10*5mm[cm]", "12*10*5 foo [cm]", "12*10*5 英寸", "cm"]:
        with pytest.raises(HTTPException):
            _unordered_specification(spec, "cm")


def test_supplier_blank_identifiers_explicit_unordered_receipt_preserves_delivery_and_payable(monkeypatch):
    import json
    from test_carton_supplier_portal import setup_portal, BASE as SUPPLIER
    with make_client(monkeypatch) as client:
        setup_portal(client)
        workbook = Workbook(); sheet = workbook.active
        sheet.append(["送货单号", "日期", "客户", "客户单号", "客户料号", "名称", "纸质", "规格", "数量", "单价"])
        sheet.append(["UNORDERED-DN-1", "2026-09-24", "华兴", "", "", "普通箱", "B3B", "10*20*30 cm", 10, 2])
        output = BytesIO(); workbook.save(output)
        upload = {"file": ("无落单送货.xlsx", output.getvalue())}
        preview = client.post(SUPPLIER + "/shipments/import-preview", files=upload)
        assert preview.status_code == 200, preview.text
        group = preview.json()["groups"][0]
        assert group["ready"] and group["rows"][0]["status"] == "AD_HOC_REVIEW"
        sent = client.post(SUPPLIER + "/shipments/import-confirm", files=upload, data={
            "sha256": preview.json()["sha256"], "selections": json.dumps([{"factory_id": "huaxing", "delivery_note_no": "UNORDERED-DN-1"}])})
        assert sent.status_code == 200, sent.text
        shipment = sent.json()["shipments"][0]
        assert shipment["lines"][0]["item_no"] == shipment["lines"][0]["contract_no"] == ""
        login_as(client, "warehouse_keeper")
        locations = client.get(BASE + "/inventory/locations", params={"factory_id": "huaxing"}).json()
        location = next(row["id"] for row in locations if row["bin_code"] == "A-01")
        line = {"shipment_line_id": shipment["lines"][0]["id"], "received_quantity": 10, "unit_price": 2,
                "location_allocations": [{"location_id": location, "quantity": 10}], "no_order_decision": "UNORDERED", "dimension_unit": "cm"}
        body = {"factory_id": "huaxing", "request_id": uuid4().hex, "expected_revision": shipment["revision"],
                "acceptance_date": "2026-09-24", "lines": [line]}
        url = SUPPLIER + f"/internal/shipments/{shipment['id']}/receive"
        assert client.post(url, json={**body, "lines": [{**line, "no_order_decision": ""}]}).status_code == 422
        assert client.post(url, json={**body, "lines": [{**line, "unit_price": 0}]}).status_code == 422
        rejected = client.post(url, json=body)
        assert rejected.status_code == 422 and "来源" in rejected.text
        assert ledger(client) == []
        latest = client.get(SUPPLIER + "/internal/workspace", params={"factory_id": "huaxing"}).json()
        assert next(row for row in latest["shipments"] if row["id"] == shipment["id"])["status"] == "SENT"
