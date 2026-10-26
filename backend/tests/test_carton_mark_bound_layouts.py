"""Reusable layouts must preserve row, face and DC provenance."""
import io
import json

import pytest
from openpyxl import Workbook
from app.services import carton_mark_flexible as flow, carton_mark_compact_plan as compact


def source():
    book=Workbook();sheet=book.active;sheet.title='Data'
    sheet.append(['Customer PO#','Customer PO item No.','ZURU PO No.','SKU No.','Description','Qty','Case Pack','DC#881','DC#884'])
    sheet.append([None,None,'PO-1','001','A',24,6,4,0])
    sheet.append([None,None,'PO-1','002','B',12,6,0,2])
    out=io.BytesIO();book.save(out);book.close()
    return flow.read('input.xlsx',out.getvalue(),allow_unresolved=True)


def plan(src):
    # Basic tests do not declare a destination matrix.
    src['cells']=[c.model_copy(update={'text':c.text.replace('DC#','Destination ')}) for c in src['cells']]
    ids={c.location.split('!')[1]:c.id for c in src['cells']}
    layout=dict(id='L1',panels=[dict(face='front',copies=2,logo=False,fields=[
        dict(label_id=ids['D1'],column_ids=[ids['D1']],role='normal'),
        dict(label_id=ids['E1'],column_ids=[ids['E1']],role='normal')],barcodes=[],symbols=[])])
    return dict(template_id='source-flow-v1',paper='A3',orientation='landscape',layouts=[layout],
                records=[dict(layout_id='L1',row_ids=[r['id'] for r in src['detail_rows']])]),ids


def test_reusable_layout_uses_populated_po_and_current_row_without_printing_unrequested_po():
    src=source();raw,ids=plan(src)
    assert len(src['detail_rows'])==2
    assert src['detail_rows'][0]['po_id']==ids['C2']
    result=flow.validate(compact.expand(json.dumps(raw),src),src,review_mode=True)
    assert [r.row_id for r in result.records]==['Data!row:2','Data!row:3']
    assert [r.panels[0].fields[0].parts[0].value for r in result.records]==['001','002']
    assert all('PO-1' not in p.value for r in result.records for f in r.panels[0].fields for p in f.parts)


@pytest.mark.parametrize('change',['unknown_row','cross_header','different_row_value','duplicate','fake_variant'])
def test_bound_layout_refuses_unsafe_row_and_column_bindings(change):
    src=source();raw,ids=plan(src)
    if change=='unknown_row':raw['records'][0]['row_ids'][0]='Data!row:99'
    if change=='cross_header':
        src['cells'].append(src['cells'][0].model_copy(update={'id':'Eforeign','location':'Data!D99','text':'Unrelated'}))
        raw['layouts'][0]['panels'][0]['fields'][0]['column_ids']=['Eforeign']
    if change=='different_row_value':raw['layouts'][0]['panels'][0]['fields'].append(dict(label_id=None,source_ids=[ids['E2']],role='normal'))
    if change=='duplicate':raw['records']*=2
    if change=='fake_variant':raw['records'][0]['variant_id']='invented'
    with pytest.raises(flow.pilot.PilotError):flow.validate(compact.expand(json.dumps(raw),src),src,review_mode=True)


def test_dc_variants_select_only_nonzero_native_rows_and_own_carton_count():
    src=source();raw,ids=plan(src)
    src['cells']=[c.model_copy(update={'text':c.text.replace('Destination ','DC#')}) for c in src['cells']]
    first=raw['layouts'][0];second=json.loads(json.dumps(first));second['id']='L2'
    first['panels'][0]['fields'].append(dict(label_id=ids['H1'],column_ids=[ids['H1']],role='normal'))
    second['panels'][0]['fields'].append(dict(label_id=ids['I1'],column_ids=[ids['I1']],role='normal'))
    raw['layouts'].append(second)
    raw['records']=[dict(layout_id=l,row_ids=['Data!row:2','Data!row:3'],when_column_id=h,variant_id=h) for l,h in [('L1',ids['H1']),('L2',ids['I1'])]]
    result=flow.validate(compact.expand(json.dumps(raw),src),src,review_mode=True)
    assert len(result.records)==2
    assert [(r.row_id,r.variant_id) for r in result.records]==[('Data!row:2',ids['H1']),('Data!row:3',ids['I1'])]
    assert [r.panels[0].fields[-1].parts[0].value for r in result.records]==['4','2']
    raw['records'].pop()
    with pytest.raises(flow.pilot.PilotError):flow.validate(compact.expand(json.dumps(raw),src),src,review_mode=True)


def test_repeated_traditional_and_stacked_face_headers_have_bounded_regions():
    book=Workbook();sheet=book.active
    for address,value in {'D3':'正嘜(FRONT MARK)X2','M3':'側嘜(SIDE MARK)X2','B8':'Front text','I8':'Side text',
                          'D15':'正嘜(FRONT MARK)X2','M15':'側嘜(SIDE MARK)X2','B20':'Second front','I20':'Second side'}.items():sheet[address]=value
    stacked=book.create_sheet('Stacked');stacked['A1']='CARTON MAIN MARK(TWO SIDES):';stacked['B3']='Importer'
    stacked['A10']='CARTON SIDE MARK(TWO SIDES):';stacked['B12']='Product'
    out=io.BytesIO();book.save(out);book.close();src=flow.read('input.xlsx',out.getvalue())
    faces={c.location:src['ownership'].get(c.id) for c in src['cells']}
    assert faces['Sheet!B8']==faces['Sheet!B20']=='front'
    assert faces['Sheet!I8']==faces['Sheet!I20']=='side'
    assert faces['Stacked!B3']=='front' and faces['Stacked!B12']=='side'
    assert src['face_copies']['Sheet']=={'front':2,'side':2}


def test_formula_reference_metadata_corrects_quantity_column_without_evaluating_formulas():
    from zipfile import ZipFile,ZIP_DEFLATED
    from xml.etree import ElementTree as ET
    b=Workbook();data=b.active;data.title='Data'
    data.append(['PO','Item No','Description','Total Quantity','Case Pack'])
    data.append(['PO-1','001','A',240,6])
    s=b.create_sheet('Print');s['A1']='Main mark (two sides)';s['A3']='Quantity:';s['B3']='=Data!E2';s['C3']='PCS'
    s['B4']='=B3*40';s['B5']="='[external.xlsx]Data'!E2"
    out=io.BytesIO();b.save(out);b.close();patched=io.BytesIO()
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with ZipFile(io.BytesIO(out.getvalue())) as original,ZipFile(patched,'w',ZIP_DEFLATED) as target:
        for name in original.namelist():
            content=original.read(name)
            if name=='xl/worksheets/sheet2.xml':
                tree=ET.fromstring(content)
                for c in tree.findall('.//m:c',ns):
                    if c.attrib['r'] in {'B3','B4','B5'}:c.find('m:v',ns).text={'B3':'6','B4':'240','B5':'6'}[c.attrib['r']]
                content=ET.tostring(tree)
            target.writestr(name,content)
    src=flow.read('formula.xlsx',patched.getvalue(),allow_unresolved=True);ids={c.location:c.id for c in src['cells']}
    assert src['native_links']=={ids['Print!B3']:ids['Data!E2']}
    raw=dict(template_id='source-flow-v1',paper='A3',orientation='landscape',records=[dict(row_id='Data!row:2',panels=[dict(face='front',copies=2,logo=False,
        fields=[dict(label_id=ids['Print!A3'],source_ids=[ids['Data!D2']],role='normal')],barcodes=[],symbols=[])])])
    valid=flow.validate(compact.expand(json.dumps(raw),src),src,review_mode=True)
    assert valid.records[0].panels[0].fields[0].parts[0].source_id==ids['Data!E2']
    assert valid.records[0].panels[0].fields[0].parts[0].value=='6'
    assert [p.value for p in valid.records[0].panels[0].fields[0].parts]==['6','PCS']
    src['native_links']={}
    raw['records'][0]['panels'][0]['fields'][0]['source_ids']=[ids['Data!E2']]
    restored=flow.validate(compact.expand(json.dumps(raw),src),src,review_mode=True)
    assert [p.value for p in restored.records[0].panels[0].fields[0].parts]==['6','PCS']


def test_native_print_blocks_require_complete_unique_po_and_barcode_matching():
    from app.services.carton_mark_bound_sources import build
    b=Workbook();s=b.active;s.title='Data'
    s.append(['Customer PO No.','Item No.','Description','Case Pack','UPC'])
    s.append(['PO-A','001','A',6,'4006381333931']);s.append(['PO-B','002','B',12,'193052037961'])
    prints=b.create_sheet('Print')
    for start,po,code,qty in [(1,'PO-A','4006381333931','6'),(20,'PO-B','193052037961','12')]:
        for address,value in {f'A{start}':'Front Marking (Four side printing)',f'A{start+2}':'PO Number:',f'B{start+2}':po,
             f'A{start+3}':'Case Pack:',f'B{start+3}':qty,f'A{start+4}':'Date of Manufacture:',f'B{start+4}':'Jun.-2026',
             f'A{start+5}':'Carton Dimension:',f'B{start+5}':'20x30x40 cm',f'A{start+6}':'UPC:',f'B{start+7}':code,
             f'A{start+8}':'MADE IN CHINA'}.items():prints[address]=value
    out=io.BytesIO();b.save(out);b.close();src=flow.read('print.xlsx',out.getvalue())
    valid=build(src);assert valid and len(valid.records)==2
    assert all(r.panels[0].copies==4 for r in valid.records)
    assert [r.row_id for r in valid.records]==['Data!row:2','Data!row:3']
    src['blocks']=src['blocks'][:1]
    assert build(src) is None  # A partial layout never becomes automatic success.


def test_linked_layout_uses_printed_sku_and_keeps_logo_only_front():
    from app.services.carton_mark_ai_pilot import Evidence
    from app.services.carton_mark_bound_sources import linked
    def cell(key,location,text):return Evidence(id=key,kind='current_cell',location=location,text=text)
    cells=[cell('E1','Data!A1','PO'),cell('E2','Data!B1','ITEM NO.'),cell('E3','Data!C1','SKU NO.'),cell('E4','Data!D1','Case Pack'),
           cell('E5','Data!A2','PO-1'),cell('E6','Data!B2','INTERNAL-99'),cell('E7','Data!C2','SKU-001'),cell('E8','Data!D2','6'),
           cell('E9','Print!F3','Item No:'),cell('E10','Print!G3','SKU-001'),cell('E11','Print!F4','Qty/Ctn:'),cell('E12','Print!G4','6')]
    src=dict(cells=cells,notes=[Evidence(id='N1',kind='current_note',location='Print!J1',text='四面ZURU LOGO')],ownership={k:'side' for k in ['E9','E10','E11','E12']},
        native_links={'E10':'E7','E12':'E8'},detail_rows=[dict(id='Data!row:2',po_id='E5',header_row=1,cell_ids={'E5','E6','E7','E8'})],
        blocks=[dict(sheet='Print',start_row=1,face='front',copies=2,cell_ids=[]),dict(sheet='Print',start_row=1,face='side',copies=2,cell_ids=['E9','E10','E11','E12'])])
    plan=linked(src);assert plan
    assert plan.records[0].panels[0].face=='front' and plan.records[0].panels[0].fields==[]
    assert plan.records[0].panels[1].fields[0].parts[0].value=='SKU-001'
    assert all(p.value!='INTERNAL-99' for pan in plan.records[0].panels for f in pan.fields for p in f.parts)


def test_unnumbered_logo_and_scanning_requirements_are_native_notes():
    b=Workbook();s=b.active;s['A1']='Front mark (two sides)';s['A3']='We will add ZURU logo in the middle of four sides.'
    s['A4']='外箱四面中间添加ZURU LOGO。';s['A5']='条码需要扫出来';s['A6']='scanable Dun 14 barcode size - 3” L x 0.75” H'
    out=io.BytesIO();b.save(out);b.close();src=flow.read('requirements.xlsx',out.getvalue(),allow_unresolved=True)
    assert len(src['notes'])==4
    assert flow.constraints(src)['logo']
    assert flow.constraints(src)['width']==pytest.approx(76.2)
    assert flow.constraints(src)['height']==pytest.approx(19.05)


def test_single_native_two_face_template_preserves_columns_units_and_each_item():
    from app.services.carton_mark_bound_sources import table_template
    b=Workbook();s=b.active;s.title='Data'
    s.append(['PO','Item No.','New Description','QTY: PCS','Carton No.:','W. G.R','W. NET','MEASUREMENT','GTIN'])
    s.append(['PO-A','001','Product A',6,'/30','4.71','3.87','63.8 x 40 x 37.8','10193052037968'])
    s.append(['PO-A','002','Product B',12,'/40','5.35','4.41','81.6 x 32.1 x 42.3','10193052037951'])
    p=b.create_sheet('Print')
    values={'A4':'正唛','F5':'侧唛','C7':'10193052037968','B8':'Item No.','C8':'001','E8':'ITEM #:','F8':'001',
            'B9':'Description','C9':'Product A','E9':'QTY:','F9':'6','G9':'PCS','B10':'QTY:','C10':'6','D10':'PCS',
            'E10':'W. G.R.:','F10':'4.71','G10':'KGS','B11':'Ctn No.:','C11':'/30','E11':'W. NET :','F11':'3.87','G11':'KGS',
            'B12':'Order No:','C12':'PO-A','E12':'MEASUREMENT:','F12':'63.8 x 40 x 37.8','G12':'CM','B13':'MADE IN CHINA'}
    for address,value in values.items():p[address]=value
    out=io.BytesIO();b.save(out);b.close();src=flow.read('labels.xlsx',out.getvalue(),allow_unresolved=True)
    plan=table_template(src);assert plan and len(plan.records)==2
    front,side=plan.records[1].panels
    assert front.face=='front' and side.face=='side'
    assert [part.value for part in front.fields[2].parts]==['12','PCS']
    assert side.fields[2].parts[0].value=='5.35'
    assert front.barcodes[0].source.value=='10193052037951' and not side.barcodes


def test_spanish_accents_use_a_latin_font_instead_of_a_missing_cjk_glyph():
    from pypdf import PdfReader
    b=Workbook();b.active['A1']='TARÁNTULA GIGANTE';out=io.BytesIO();b.save(out);b.close()
    src=flow.read('spanish.xlsx',out.getvalue());value=src['cells'][0]
    plan=flow.validate(json.dumps(dict(template_id='source-flow-v1',paper='A3',orientation='landscape',records=[dict(panels=[dict(face='front',copies=1,logo=False,fields=[dict(label=None,parts=[dict(source_id=value.id,value=value.text)],role='normal')],barcodes=[],symbols=[])])])),src)
    pdf,_,_=flow.render(plan,src,include_sources=False);fonts=[]
    PdfReader(io.BytesIO(pdf)).pages[0].extract_text(visitor_text=lambda text,cm,tm,font,size:fonts.append(font.get('/BaseFont')) if 'TARÁNTULA' in text and font else None)
    assert fonts==['/Helvetica']
