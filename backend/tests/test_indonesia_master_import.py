"""Material master import is explicit, atomic, scoped, and separate from shipment facts."""
import asyncio
from copy import deepcopy
from io import BytesIO
import json
from uuid import uuid4

import openpyxl
import pytest
from fastapi import HTTPException

from test_indonesia_materials import client, actor, command, material, product, order, workspace, supply, allocation, shipment, BASE
from app.services.indonesia_master_import import parse_material_book
from app.services.customer_order_jobs import run_order_job
from app.schemas.indonesia_materials import MasterSheetMapping

HEADERS = ['货号', '物料编码', '产品中文名称\nChinese name of the product', '规格', '数量单位', '报关单位', '单价', '币种', '单件用量', '送货数量', '单个净重（kg）']


def workbook(rows, headers=HEADERS, merged=False):
    book = openpyxl.Workbook(); sheet = book.active; sheet.title = '出货明细'
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    if merged:
        sheet.merge_cells('A2:A3')
    stream = BytesIO(); book.save(stream)
    return stream.getvalue()


def source(code='M-NEW', name='正式报关名称', item='40149', **changes):
    data = dict(zip(HEADERS, [item, code, name, '规格A', 'pcs', 'PCE', '0.5', 'USD', '2', '8000', '.01']))
    data.update(changes)
    return [data.get(key) for key in HEADERS]


def catalog(client):
    m = material(client, code='EXISTING')
    p = product(client, m, code='40149')
    other = product(client, m, code='40379', materials=[])
    unrelated = product(client, m, code='99999')
    return m, p, other, unrelated


def preview(client, content, *products, status=200, **context):
    response = client.post(BASE + '/master-imports/preview', data=dict(factory_id='huaxing', customer='合成客户',
        context=json.dumps(dict(product_ids=[p['id'] for p in products]) | context)), files={'file': ('统一资料.xlsx', content)})
    assert response.status_code == status, response.text
    return response.json()['record'] if status == 200 else response


def choices(record, decision='import'):
    result = []
    for g in record['data']['groups']:
        row = g['sources'][0]
        incoming = {key: value for key, value in row['source'].items() if key not in {'item_no', 'piece_usage'}}
        current = g['existing']['data'] if g['existing'] else {}
        if current:
            incoming['code'] = current['code']
        result.append(dict(group=g['index'], source_index=row['index'], decision=decision,
            material=(current | incoming) if decision == 'import' else None,
            recipes=[dict(product_id=p['id'], decision='import', usage=row['source'].get('piece_usage') or p['existing_usage']) for p in g['products']]))
    return dict(choices=result)


def test_selected_products_only_new_material_and_additive_recipes(client):
    m, p, q, outside = catalog(client)
    before = workspace(client)
    content = workbook([source(item='40149 / 40379'), source(code='OUTSIDE', item='99999'), source(code='PREFIX', item='401490')])
    staged = preview(client, content, p, q)
    assert staged['data']['matched_rows'] == 1 and staged['data']['skipped_rows'] == 2
    assert len(workspace(client)['records']) == len(before['records']) + 1  # preview only, no masters
    result = command(client, 'master_import_confirm', choices(staged), staged)
    records = {r['id']: r for r in workspace(client)['records']}
    mid = result['data']['choices'][0]['material_id']
    assert records[mid]['data']['name'] == '正式报关名称'
    assert records[mid]['data']['billing_unit'] == 'PCE'
    assert records[p['id']]['data']['materials'] == p['data']['materials'] + [dict(material_id=mid, usage='2')]
    assert records[q['id']]['data']['materials'] == [dict(material_id=mid, usage='2')]
    assert records[outside['id']] == outside
    assert not any(r['kind'] in {'supply', 'order', 'shipment'} for r in records.values())


def test_duplicates_compare_explicit_keep_or_import_and_snapshots_remain(client):
    m, p, q, outside = catalog(client)
    o = order(client, item_no='40149')
    original = deepcopy(o)
    staged = preview(client, workbook([source('EXISTING', m['data']['name'], **{'规格': 'TEST'}), source('EXISTING', m['data']['name'], **{'规格': 'TEST', '单价': '.7'})]), p)
    group = staged['data']['groups'][0]
    assert len(group['sources']) == 2 and group['existing'] == m
    assert {r['code'] for r in group['affected_products']} == {'40149', '99999'}
    payload = choices(staged, 'keep'); payload['choices'][0]['recipes'][0].update(decision='keep', usage=None)
    command(client, 'master_import_confirm', payload, staged)
    rows = {r['id']: r for r in workspace(client)['records']}
    assert rows[m['id']] == m and rows[p['id']] == p
    fresh = preview(client, workbook([source('EXISTING', m['data']['name'], **{'规格': 'TEST'})]), p)
    command(client, 'master_import_confirm', choices(fresh), fresh)
    rows = {r['id']: r for r in workspace(client)['records']}
    assert rows[m['id']]['data']['price'] == '0.5'
    assert rows[m['id']]['data']['image_data'] == m['data']['image_data']
    assert rows[o['id']] == original


def test_missing_usage_never_uses_shipment_quantity_atomic_rollback(client):
    _, p, q, _ = catalog(client)
    staged = preview(client, workbook([source(code='OK'), source(code='NO-USAGE', name='另一物料', **{'单件用量': None})]), p, q)
    assert staged['data']['groups'][1]['sources'][0]['source'].get('piece_usage') is None
    before = deepcopy(workspace(client)['records'])
    command(client, 'master_import_confirm', choices(staged), staged, status=422)
    assert workspace(client)['records'] == before


@pytest.mark.parametrize('target', ['product', 'material', 'new-code', 'shared-reference'])
def test_concurrent_changes_reject_without_partial_write(client, target):
    m, p, q, outside = catalog(client)
    code = 'NEW' if target == 'new-code' else 'EXISTING'
    incoming = source(code, m['data']['name'], **{'规格': 'TEST'})
    staged = preview(client, workbook([incoming]), p)
    if target == 'product':
        command(client, 'product_save', p['data'] | {'name': '并发'}, p)
    elif target == 'material':
        command(client, 'material_save', m['data'] | {'name': '并发'}, m)
    elif target == 'new-code':
        material(client, code='NEW')
    else:
        command(client, 'product_save', q['data'] | {'materials': [dict(material_id=m['id'], usage='1')]}, q)
    before = deepcopy(workspace(client)['records'])
    command(client, 'master_import_confirm', choices(staged), staged, status=409)
    assert workspace(client)['records'] == before
    if target == 'shared-reference':
        fresh = preview(client, workbook([incoming]), p)
        assert fresh['id'] != staged['id']
        command(client, 'master_import_confirm', choices(fresh), fresh)


def test_no_code_matching_preserves_significant_specification_punctuation(client):
    m = material(client, code='S1', name='塑胶', specification='10*20')
    p = product(client, m, code='40149')
    staged = preview(client, workbook([source('', '塑胶', **{'规格': '1020'}), source('', '塑胶', **{'规格': '10*20'})]), p)
    assert staged['data']['groups'][0]['existing'] is None
    assert staged['data']['groups'][1]['existing']['id'] == m['id']


def test_ambiguous_name_and_archived_code_blocked(client):
    m = material(client, code='A'); material(client, code='B')
    p = product(client, m, code='40149')
    staged = preview(client, workbook([source('', m['data']['name'], **{'规格': 'TEST'})]), p)
    assert '多条' in staged['data']['groups'][0]['blocked']
    command(client, 'master_import_confirm', choices(staged), staged, status=422)
    command(client, 'product_save', p['data'] | {'materials': [], 'material_deletions': [dict(material_id=m['id'], expected_version=m['version'])]}, p)
    fresh_p = next(r for r in workspace(client)['records'] if r['id'] == p['id'])
    staged = preview(client, workbook([source('A', m['data']['name'], **{'规格': 'TEST'})]), fresh_p)
    assert '删除' in staged['data']['groups'][0]['blocked']


def test_explicit_choice_required_and_retries_are_idempotent(client):
    _, p, _, _ = catalog(client)
    staged = preview(client, workbook([source()]), p)
    payload = choices(staged)
    bad = deepcopy(payload); bad['choices'][0].pop('decision')
    command(client, 'master_import_confirm', bad, staged, status=422)
    command(client, 'import_confirm', {'indexes': [0]}, staged, status=422)
    operation = uuid4().hex
    first = command(client, 'master_import_confirm', payload, staged, operation_id=operation)
    assert command(client, 'master_import_confirm', payload, staged, operation_id=operation) == first
    command(client, 'master_import_confirm', payload, staged, status=409)


def test_skipped_groups_can_be_imported_on_later_preview(client):
    m, p, _, _ = catalog(client)
    content = workbook([source('EXISTING', m['data']['name'], **{'规格': 'TEST'}), source('LATER')])
    staged = preview(client, content, p)
    payload = choices(staged, 'keep'); payload['choices'] = payload['choices'][:1]
    payload['choices'][0]['recipes'][0].update(decision='keep', usage=None)
    command(client, 'master_import_confirm', payload, staged)
    fresh = preview(client, content, p)
    assert fresh['id'] != staged['id'] and fresh['data']['status'] == 'preview'
    payload = choices(fresh); payload['choices'] = payload['choices'][1:]
    command(client, 'master_import_confirm', payload, fresh)


def test_factory_customer_and_permission_boundaries(client):
    _, p, _, _ = catalog(client)
    content = workbook([source()])
    other = dict(p, id='another-factory-product')
    preview(client, content, other, status=404)
    client.materials_auth['user'] = actor('read')
    preview(client, content, p, status=403)
    client.materials_auth['user'] = actor('read', 'write', factory='huakang')
    preview(client, content, p, status=403)


def test_parser_preserves_formatted_codes_merged_products_and_ignores_irrelevant_formula_columns():
    content = workbook([source(code='000001'), source(code='000002', item=None)], merged=True)
    parsed = parse_material_book(content, '统一.xlsx')
    assert [r['source']['item_no'] for r in parsed['rows']] == ['40149', '40149']
    assert parsed['rows'][0]['source']['code'] == '000001'
    assert 'quantity' not in parsed['rows'][0]['source']
    formula = parse_material_book(workbook([source(name='=A1', code=''), source(item='=40149'), source(**{'单个净重（kg）': '=1/1000'})]), '统一.xlsx')
    assert len(formula['issues']) == 2
    assert len(formula['rows']) == 1 and formula['rows'][0]['warnings']
    assert 'net_weight_kg' not in formula['rows'][0]['source']


def test_manual_header_mapping_units_and_source_column_validation():
    headers = ['货号\nItem No.', '产品中文名称\nChinese name of the product', '单位\nUnit', '单位\nUnit', '数量', '用量']
    content = workbook([['40149', '正式名称', 'KGM', 'pcs', '1234', '2']], headers)
    automatic = parse_material_book(content, '统一.xlsx')
    assert automatic['sheets'][0]['columns'] == {'item_no': 1, 'name': 2}
    mapping = MasterSheetMapping(sheet='出货明细', header_row=1, columns={'item_no': 1, 'name': 2, 'customs_unit': 3, 'unit': 4, 'piece_usage': 6})
    parsed = parse_material_book(content, '统一.xlsx', [mapping])
    assert parsed['rows'][0]['source']['customs_unit'] == 'KGM'
    assert parsed['rows'][0]['source']['unit'] == 'pcs'
    mapping.columns['piece_usage'] = 4
    with pytest.raises(HTTPException):
        parse_material_book(content, '统一.xlsx', [mapping])


def test_real_bounded_worker_can_parse_material_book():
    result = asyncio.run(run_order_job(parse_material_book, workbook([source()]), '统一.xlsx', []))
    assert result['rows'][0]['source']['code'] == 'M-NEW'


def test_imported_recipe_continues_through_demand_receipt_allocation_and_excel(client):
    original, _, p, _ = catalog(client)
    staged = preview(client, workbook([source(item='40379')]), p)
    payload = choices(staged)
    # Business supplements export-only fields absent from the source workbook.
    payload['choices'][0]['material'].update({key: original['data'][key] for key in
        ('english_name', 'category', 'usage', 'department', 'export_company', 'image_data')})
    result = command(client, 'master_import_confirm', payload, staged)
    mid = result['data']['choices'][0]['material_id']
    m = next(r for r in workspace(client)['records'] if r['id'] == mid)
    o = order(client, item_no='40379', quantity='10')
    assert o['data']['demands'][0]['quantity'] == '20'
    received = command(client, 'receipt_confirm', {'actual_date': '2026-10-28'},
                       supply(client, m, normal='20', spare='0', purchase_quantity='20'))
    assigned = allocation(client, received, o, normal='20', spare='0')
    draft = shipment(client, assigned, normal='20', spare='0', shipment_date='2026-10-28')
    confirmed = command(client, 'shipment_confirm', {}, draft)
    response = client.get(BASE + '/shipments/' + confirmed['id'] + '/export', params={'factory_id': 'huaxing'})
    assert response.status_code == 200
    book = openpyxl.load_workbook(BytesIO(response.content), data_only=False)
    detail = book['5.21出货明细']
    assert '正式报关名称' in [cell.value for row in detail.iter_rows() for cell in row]
    assert len(book.sheetnames) > 1


@pytest.mark.parametrize('shared_code', ['53003116', ''])
def test_shared_code_four_names_import_receipt_match_and_mixed_export_remain_separate(client, shared_code):
    from test_indonesia_warehouse_import import preview as receipt_preview, source as receipt_source, choose, confirm as receipt_confirm
    names = ['塑胶件-子弹筒', '塑胶件-绿色拉杆', '塑胶件-蓝色拉杆', '塑胶件-推杆']
    m = material(client, code=shared_code, name=names[0])
    p = product(client, m, code='40149')
    staged = preview(client, workbook([source(shared_code, name, **{'规格': 'TEST'}) for name in names]), p)
    assert len(staged['data']['groups']) == 4
    assert all(bool(g['code_warnings']) == bool(shared_code) for g in staged['data']['groups'])
    payload = choices(staged)
    for choice in payload['choices']:
        choice['material'] = m['data'] | choice['material']
    saved = command(client, 'master_import_confirm', payload, staged)
    ids = [choice['material_id'] for choice in saved['data']['choices']]
    assert len(set(ids)) == 4 and ids[0] == m['id']
    o = order(client, item_no='40149', quantity='10')
    assert len(o['data']['demands']) == 4
    receipts = receipt_preview(client, *[receipt_source(material_code=shared_code, material_name=name,
        receipt_line=str(index + 1), purchase_line=str(index + 1), normal_quantity='20', spare_quantity='0') for index, name in enumerate(names)])
    assert [row['auto_material_id'] for row in receipts['data']['rows']] == ids
    ambiguous = receipt_preview(client, receipt_source(receipt_no='WH-AMB', material_code=shared_code, material_name=''))
    assert ambiguous['data']['rows'][0]['auto_material_id'] is None
    if shared_code:
        assert len(ambiguous['data']['rows'][0]['materials']) == 4
    else:
        assert ambiguous['data']['rows'][0]['blocked']
    posted = receipt_confirm(client, receipts, selections=[choose(receipts, row=index) for index in range(4)])
    draft = next(r for r in workspace(client)['records'] if r['kind'] == 'shipment')
    package = dict(label='四种共码物料混箱', cartons=1, length_cm='30', width_cm='20', height_cm='10', carton_weight_kg='.1', items=draft['data']['lines'])
    draft = command(client, 'shipment_save', {k: draft['data'][k] for k in ('customer', 'recipient', 'shipment_date', 'lines', 'remarks')} | {'packages': [package]}, draft)
    confirmed = command(client, 'shipment_confirm', {}, draft)
    assert confirmed['data']['export_material_count'] == 4
    assert confirmed['data']['supplier_summary'][0]['materials'] == 4
    assert confirmed['data']['packing_totals']['cartons'] == 1
    response = client.get(BASE + '/shipments/' + confirmed['id'] + '/export', params={'factory_id': 'huaxing'})
    assert response.status_code == 200
    book = openpyxl.load_workbook(BytesIO(response.content), data_only=True)
    detail = book['5.21出货明细']
    exported = [row for row in detail.iter_rows() if row[5].value in names]
    assert len(exported) == 4
    assert {row[5].value for row in exported} == set(names)
    assert all(str(row[51].value or '') == shared_code and row[11].value == 20 for row in exported)


def test_import_can_fill_blank_code_only_after_explicit_adoption(client):
    m = material(client, code=''); p = product(client, m, code='40149')
    content = workbook([source('NEW-CODE', m['data']['name'], **{'规格': 'TEST'})])
    staged = preview(client, content, p)
    assert staged['data']['groups'][0]['existing']['id'] == m['id']
    assert not staged['data']['groups'][0]['code_warnings']
    command(client, 'master_import_confirm', choices(staged, 'keep'), staged)
    assert next(r for r in workspace(client)['records'] if r['id'] == m['id'])['data']['code'] == ''
    fresh = preview(client, content, p)
    payload = choices(fresh)
    payload['choices'][0]['material']['code'] = 'NEW-CODE'
    command(client, 'master_import_confirm', payload, fresh)
    saved = next(r for r in workspace(client)['records'] if r['id'] == m['id'])
    assert saved['data']['code'] == 'NEW-CODE'


def test_shared_code_new_name_allowed_but_same_named_identity_and_archives_are_protected(client):
    first = material(client, code='SHARED', name='子弹筒')
    second = material(client, code='SHARED', name='拉杆')
    assert first['id'] != second['id']
    command(client, 'material_save', second['data'], status=409)
    p = product(client, first, code='40149')
    staged = preview(client, workbook([source('SHARED', '推杆', **{'规格': 'TEST'})]), p)
    material(client, code='OTHER', name='推杆')
    before = workspace(client)['records']
    command(client, 'master_import_confirm', choices(staged), staged, status=409)
    assert workspace(client)['records'] == before


def test_old_code_grouped_preview_requires_reread(client):
    from sqlalchemy.orm import Session
    from app.models.indonesia_materials import IndonesiaMaterialRecord
    _, p, _, _ = catalog(client)
    staged = preview(client, workbook([source()]), p)
    with Session(client.engine) as db:
        row = db.get(IndonesiaMaterialRecord, staged['id'])
        data = dict(row.data); data.pop('identity_rule')
        row.data = data; db.commit()
    response = command(client, 'master_import_confirm', choices(staged), staged, status=409)
    assert '重新读取' in response.json()['detail']


def test_legacy_same_name_codes_keep_their_ids_and_allow_updates_and_archived_peer(client):
    from sqlalchemy.orm import Session
    from app.models.indonesia_materials import IndonesiaMaterialRecord
    from app.services.indonesia_materials import digest
    a = material(client, code='A'); b = material(client, code='B')
    p = product(client, a, code='40149')
    with Session(client.engine) as db:
        db.get(IndonesiaMaterialRecord, a['id']).business_key = digest('A')
        db.commit()
    staged = preview(client, workbook([source('A', a['data']['name'], **{'规格': 'TEST'}), source('B', b['data']['name'], **{'规格': 'TEST'})]), p)
    assert len(staged['data']['groups']) == 2
    saved = command(client, 'master_import_confirm', choices(staged), staged)
    assert {x['material_id'] for x in saved['data']['choices']} == {a['id'], b['id']}
    with Session(client.engine) as db:
        first = db.get(IndonesiaMaterialRecord, a['id'])
        assert first.business_key == digest('A')
        first.data = first.data | {'deleted': True}; db.commit()
    current = next(r for r in workspace(client)['records'] if r['id'] == b['id'])
    command(client, 'material_save', current['data'] | {'price': '.9'}, current)
    command(client, 'material_save', current['data'] | {'code': 'NEW'}, status=409)
