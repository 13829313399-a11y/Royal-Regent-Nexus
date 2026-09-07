from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
import pytesseract
from openpyxl.utils import get_column_letter

from app.api.customer_order import CUSTOMER_FACTORY_IDS
from app.services import customer_order_huadeng as service
from app.services.huadeng_order_legacy import common_pdf_text, simba_po_parser


def _workbook_bytes(sheet_name: str, headers: dict[int, str], row: dict[int, object]) -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    for column, value in headers.items():
        worksheet.cell(1, column, value)
    for column, value in row.items():
        worksheet.cell(2, column, value)
    for column in range(1, max(headers, default=1) + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 16
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_huadeng_mapping_registry_is_factory_scoped() -> None:
    assert set(service.HUADENG_CUSTOMER_MAPPINGS) == {
        "casdon", "jakks", "simba", "spin", "spin-master", "goliath",
    }
    assert {
        code: CUSTOMER_FACTORY_IDS[code]
        for code in service.HUADENG_CUSTOMER_MAPPINGS
    } == {
        "casdon": "huadeng",
        "jakks": "huadeng",
        "simba": "huadeng",
        "spin": "huadeng",
        "spin-master": "huadeng",
        "goliath": "huadeng",
    }
    with pytest.raises(service.HuadengCustomerOrderError, match="只属于华登厂区"):
        service.create_huadeng_customer_preview(
            customer_code="casdon",
            factory_id="huaxing",
            received_date="2026-08-03",
            po_files=[("casdon.pdf", b"pdf")],
            schedule_file_name="Casdon排期.xlsx",
            schedule_content=b"schedule",
        )


def test_revision_dedupe_prefers_latest_and_more_complete_order() -> None:
    original = {
        "po_number": "PO-1",
        "filename": "PO-1.pdf",
        "po_date": "2026-07-01",
        "lines": [{"item_code": "A1", "qty": 10}],
    }
    revision = {
        "po_number": "PO-1",
        "filename": "PO-1 REV2.pdf",
        "po_date": "2026-07-02",
        "ship_date": "2026-08-10",
        "lines": [{"item_code": "A1", "qty": 20, "unit_price": 1, "total_usd": 20}],
    }
    kept, warnings = service._dedupe_revision_orders(
        [original, revision], quality_fields=("po_date", "ship_date"),
    )
    assert kept == [revision]
    assert "保留 PO-1 REV2.pdf" in warnings[0]


def test_casdon_preview_and_export_use_received_date_and_seven_day_rule(monkeypatch) -> None:
    headers = {
        2: "来单日期", 4: "合同联系人", 5: "客户PO", 6: "合同号", 7: "客名",
        8: "版本", 9: "货号", 10: "产品名称", 11: "数量", 13: "外装箱",
        27: "验货期", 29: "PO走货期", 32: "单价/港币", 33: "单价/美金",
        34: "总货价/港币", 35: "总货价/美金", 37: "英文品名", 47: "走货国家",
    }
    schedule = _workbook_bytes(
        "Casdon 排货表-总",
        headers,
        {
            6: "OLD", 7: "CASDON UK", 8: "A", 9: "1234(A)",
            10: "玩具厨房", 11: 12, 13: 8, 37: "Toy Kitchen",
        },
    )
    schedule_book = openpyxl.load_workbook(BytesIO(schedule))
    schedule_sheet = schedule_book["Casdon 排货表-总"]
    schedule_sheet.cell(3, 6, "OTHER")
    schedule_sheet.cell(3, 7, "CASDON UK")
    schedule_sheet.cell(3, 9, "9999")
    schedule_sheet.cell(3, 10, "其他产品")
    schedule_sheet.cell(3, 11, 6)
    schedule_buffer = BytesIO()
    schedule_book.save(schedule_buffer)
    schedule_book.close()
    schedule = schedule_buffer.getvalue()

    def parse(_self, _path: str):
        return {
            "po_number": "8395",
            "po_date": "2026-07-30",
            "ship_date": "2026-08-10",
            "customer": "KB SALES",
            "customer_po_header": "UK-PO-1",
            "version": "A",
            "lines": [{
                "line_no": "1", "item_code": "1234", "description_en": "Toy Kitchen",
                "qty": 24, "unit": "PCS", "unit_price": 10, "total_usd": 240,
                "is_charge": False,
            }],
            "raw_text": "",
        }

    monkeypatch.setattr(service.casdon_po_parser.CasdonPOParser, "parse", parse)
    preview = service.create_huadeng_customer_preview(
        customer_code="casdon",
        factory_id="huadeng",
        received_date="2026-08-03",
        po_files=[("PO 8395.pdf", b"pdf")],
        schedule_file_name="2026年Casdon排期.xlsx",
        schedule_content=schedule,
    )
    row = preview["rows"][0]
    assert row["received_date"] == "2026-08-03"
    assert row["contract_no"] == "8395"
    assert row["product_name_zh"] == "玩具厨房"
    assert row["line_q"] == "2026-08-03"
    assert row["requested_ship_date"] == "2026-08-10"
    assert row["unit_price_hkd"] == "77.5"
    assert row["amount_hkd"] == "1860"
    assert row["units_per_carton"] == "8"
    assert row["carton_count"] == "3"

    output, file_name, exported_preview = service.export_huadeng_customer_schedule(
        customer_code="casdon",
        factory_id="huadeng",
        received_date="2026-08-03",
        po_files=[("PO 8395.pdf", b"pdf")],
        schedule_file_name="2026年Casdon排期.xlsx",
        schedule_content=schedule,
    )
    workbook = openpyxl.load_workbook(BytesIO(output), data_only=False)
    try:
        assert workbook.sheetnames == ["Casdon 排货表-总"]
        assert workbook["Casdon 排货表-总"].cell(2, 6).value == "OLD"
        assert workbook["Casdon 排货表-总"].cell(3, 6).value == "8395"
        assert workbook["Casdon 排货表-总"].cell(3, 2).value == datetime(2026, 8, 3)
        assert workbook["Casdon 排货表-总"].cell(4, 6).value == "OTHER"
        assert workbook["Casdon 排货表-总"].cell(4, 9).value == "9999"
        assert workbook["Casdon 排货表-总"].cell(3, 13).value == 8
        assert workbook["Casdon 排货表-总"].cell(3, 14).value == '=IF(M3=0,"",K3/M3)'
    finally:
        workbook.close()
    assert file_name == preview["output_file_name"]
    assert exported_preview["summary"]["total"] == 1


def test_casdon_export_inserts_before_matching_item_subtotal(monkeypatch) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Casdon 排货表-总"
    headers = {
        2: "来单日期", 4: "合同联系人", 5: "客户PO", 6: "合同号", 7: "客名",
        8: "版本", 9: "货号", 10: "产品名称", 11: "数量", 13: "外装箱",
        14: "总箱", 27: "验货期", 29: "PO走货期", 32: "单价/港币",
        33: "单价/美金", 34: "总货价/港币", 35: "总货价/美金",
        37: "英文品名", 47: "走货国家",
    }
    for column, value in headers.items():
        worksheet.cell(1, column, value)
    existing_rows = (
        {6: "OLD-1", 7: "CASDON UK", 8: "A", 9: "31650(A)", 10: "玩具厨房", 11: 100, 13: 8},
        {6: "OLD-2", 7: "CASDON UK", 8: "A", 9: "31650(A)", 10: "玩具厨房", 11: 200, 13: 8},
        {10: "合计：", 11: "=SUM(K2:K3)", 14: "=SUM(N2:N3)"},
        {6: "OLD-3", 7: "CASDON UK", 8: "A", 9: "32050(A)", 10: "玩具厨房", 11: 300, 13: 4},
        {10: "合计：", 11: "=SUM(K5:K5)", 14: "=SUM(N5:N5)"},
    )
    for row_no, values in enumerate(existing_rows, start=2):
        for column, value in values.items():
            worksheet.cell(row_no, column, value)
    notes = workbook.create_sheet("说明")
    notes["A1"] = "保留工作表"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    schedule = output.getvalue()

    def parse(_self, _path: str):
        return {
            "po_number": "NEW-31650",
            "po_date": "2026-08-01",
            "ship_date": "2026-09-01",
            "customer": "CASDON UK",
            "customer_po_header": "UK-PO-31650",
            "version": "A",
            "lines": [
                {
                    "line_no": "1", "item_code": "31650", "description_en": "Toy Kitchen",
                    "qty": 24, "unit": "PCS", "unit_price": 2, "total_usd": 48,
                    "is_charge": False,
                },
                {
                    "line_no": "2", "item_code": "32050", "description_en": "Toy Kitchen",
                    "qty": 12, "unit": "PCS", "unit_price": 2, "total_usd": 24,
                    "is_charge": False,
                },
            ],
            "raw_text": "",
        }

    monkeypatch.setattr(service.casdon_po_parser.CasdonPOParser, "parse", parse)
    exported, _file_name, preview = service.export_huadeng_customer_schedule(
        customer_code="casdon",
        factory_id="huadeng",
        received_date="2026-08-03",
        po_files=[("PO NEW-31650.pdf", b"pdf")],
        schedule_file_name="2026年Casdon排期.xlsx",
        schedule_content=schedule,
    )

    result = openpyxl.load_workbook(BytesIO(exported), data_only=False)
    try:
        sheet = result["Casdon 排货表-总"]
        assert result.sheetnames == ["Casdon 排货表-总", "说明"]
        assert sheet.cell(4, 6).value == "NEW-31650"
        assert sheet.cell(4, 9).value == "31650(A)"
        assert sheet.cell(4, 13).value == 8
        assert sheet.cell(4, 14).value == '=IF(M4=0,"",K4/M4)'
        assert sheet.cell(5, 10).value == "合计："
        assert sheet.cell(5, 11).value == "=SUM(K2:K4)"
        assert sheet.cell(5, 14).value == "=SUM(N2:N4)"
        assert sheet.cell(6, 9).value == "32050(A)"
        assert sheet.cell(7, 6).value == "NEW-31650"
        assert sheet.cell(7, 9).value == "32050(A)"
        assert sheet.cell(7, 13).value == 4
        assert sheet.cell(7, 14).value == '=IF(M7=0,"",K7/M7)'
        assert sheet.cell(8, 11).value == "=SUM(K6:K7)"
        assert result["说明"]["A1"].value == "保留工作表"
    finally:
        result.close()
    assert preview["rows"][0]["units_per_carton"] == "8"
    assert preview["rows"][0]["carton_count"] == "3"
    assert preview["rows"][1]["units_per_carton"] == "4"
    assert preview["rows"][1]["carton_count"] == "3"
    assert any("插入该货号最后一条历史明细之后" in warning for warning in preview["warnings"])


def test_casdon_outer_pack_conflict_or_absence_can_be_manually_resolved(monkeypatch) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "2026年未验货订单"
    headers = {
        2: "来单日期", 4: "合同联系人", 5: "客户PO", 6: "合同号", 7: "客名",
        8: "版本", 9: "货号", 10: "产品名称", 11: "数量", 13: "外装箱",
        14: "总箱", 27: "验货期", 29: "PO走货期", 32: "单价/港币",
        33: "单价/美金", 34: "总货价/港币", 35: "总货价/美金",
        37: "英文品名", 47: "走货国家",
    }
    for column, value in headers.items():
        worksheet.cell(1, column, value)
    existing_rows = (
        {6: "OLD-USA", 7: "CASDON USA", 8: "A", 9: "11050(A)", 10: "充棉机", 11: 60, 13: 6},
        {6: "OLD-TARGET", 7: "TARGET USA", 8: "A", 9: "11050.TAR001(A)", 10: "充棉机", 11: 40, 13: 4},
        {10: "合计：", 11: "=SUM(K2:K3)", 14: "=SUM(N2:N3)"},
    )
    for row_no, values in enumerate(existing_rows, start=2):
        for column, value in values.items():
            worksheet.cell(row_no, column, value)
    notes = workbook.create_sheet("说明")
    notes["A1"] = "保留工作表"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    schedule = output.getvalue()

    def parse(_self, _path: str):
        return {
            "po_number": "NEW-MANUAL-OUTER",
            "po_date": "2026-08-01",
            "ship_date": "2026-09-01",
            "customer": "KB SALES",
            "customer_po_header": "KB-PO-1",
            "version": "A",
            "lines": [
                {
                    "line_no": "1", "item_code": "11050", "description_en": "Stuffing Machine",
                    "qty": 12, "unit": "PCS", "unit_price": 2, "total_usd": 24,
                    "is_charge": False,
                },
                {
                    "line_no": "2", "item_code": "99999", "description_en": "New Item",
                    "qty": 24, "unit": "PCS", "unit_price": 1, "total_usd": 24,
                    "is_charge": False,
                },
            ],
            "raw_text": "",
        }

    monkeypatch.setattr(service.casdon_po_parser.CasdonPOParser, "parse", parse)
    preview = service.create_huadeng_customer_preview(
        customer_code="casdon",
        factory_id="huadeng",
        received_date="2026-08-03",
        po_files=[("PO NEW-MANUAL-OUTER.pdf", b"pdf")],
        schedule_file_name="2026年Casdon排期.xlsx",
        schedule_content=schedule,
    )

    assert preview["summary"]["blocked"] == 2
    assert preview["rows"][0]["issues"][0]["code"] == "ambiguous_units_per_carton"
    assert preview["rows"][0]["issues"][0]["field"] == "units_per_carton"
    assert preview["rows"][1]["issues"][0]["code"] == "missing_units_per_carton"
    issue_keys = {
        issue["skip_key"]
        for row in preview["rows"]
        for issue in row["issues"]
    }
    overrides = [
        {"row_id": "casdon-1", "field": "units_per_carton", "value": "6"},
        {"row_id": "casdon-2", "field": "units_per_carton", "value": "12"},
    ]
    exported, _file_name, _exported_preview = service.export_huadeng_customer_schedule(
        customer_code="casdon",
        factory_id="huadeng",
        received_date="2026-08-03",
        po_files=[("PO NEW-MANUAL-OUTER.pdf", b"pdf")],
        schedule_file_name="2026年Casdon排期.xlsx",
        schedule_content=schedule,
        skipped_issue_keys=issue_keys,
        manual_overrides=overrides,
    )

    result = openpyxl.load_workbook(BytesIO(exported), data_only=False)
    try:
        assert result.sheetnames == ["2026年未验货订单", "说明"]
        sheet = result["2026年未验货订单"]
        exported_rows = {
            sheet.cell(row_no, 9).value: row_no
            for row_no in range(2, sheet.max_row + 1)
            if sheet.cell(row_no, 6).value == "NEW-MANUAL-OUTER"
        }
        conflict_row = exported_rows["11050(A)"]
        missing_row = exported_rows["99999(A)"]
        assert sheet.cell(conflict_row, 13).value == 6
        assert sheet.cell(conflict_row, 14).value == f'=IF(M{conflict_row}=0,"",K{conflict_row}/M{conflict_row})'
        assert sheet.cell(missing_row, 13).value == 12
        assert sheet.cell(missing_row, 14).value == f'=IF(M{missing_row}=0,"",K{missing_row}/M{missing_row})'
        assert result["说明"]["A1"].value == "保留工作表"
    finally:
        result.close()


def test_jakks_blocks_change_orders_and_dedupes_batch(monkeypatch) -> None:
    monkeypatch.setattr(
        service.jakks_schedule,
        "load_dataset",
        lambda _path: {"records": [{"sku": "J-1", "product": "标准公仔", "contact": "Amy", "contract_no": "C1"}]},
    )

    def parse_po(_path):
        return {
            "contract_no": "C1",
            "customer_po": "PO-1",
            "customer": "JAKKS",
            "order_date": "2026-08-01",
            "lines": [{
                "item_no": "J-1", "product_name": "PO品名", "po_description": "Toy Figure",
                "quantity": 100, "unit": "PCS", "unit_price_usd": 2, "total_usd": 200,
                "ship_date": "2026-09-01", "special_note": "",
            }],
            "warnings": [],
        }

    monkeypatch.setattr(service.jakks_po_parser, "parse_po", parse_po)
    prepared = service._prepare_jakks(
        [("C1-CXL.pdf", b"cancel"), ("C1.pdf", b"one"), ("C1-copy.pdf", b"two")],
        "2026年Jakks排期.xlsx",
        b"schedule",
    )
    assert len(prepared.records) == 1
    assert prepared.records[0]["product_name"] == "标准公仔"
    assert prepared.records[0]["contact"] == "Amy"
    assert any("CXL/SUP" in warning for warning in prepared.warnings)
    assert any("完全重复" in warning for warning in prepared.warnings)


def test_jakks_standard_contract_uses_notes_not_order_notes_remarks() -> None:
    rows = [
        ["JAKKS-10001 UPC 123 Country of Origin CN", "Toy Figure Customer Item No: 88", 2.0, "100 CA Outer Pack: 10", "USD 200.00"],
        ["ORDER NOTES: REMARKS do not write this"],
        ["NOTES: Use customer-approved label"],
    ]

    parsed = service.jakks_po_parser._parse_standard_contract_rows(rows)

    assert len(parsed) == 1
    assert parsed[0]["special_note"] == "Use customer-approved label"


def test_jakks_scanned_pdf_prefers_rapidocr(monkeypatch) -> None:
    from app.services import carton_mark

    class FakePage:
        def extract_text(self, **_kwargs):
            return ""

        def to_image(self, **kwargs):
            assert kwargs["resolution"] == 180
            return type("Rendered", (), {"original": object()})()

    class FakePdf:
        pages = [FakePage()]

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

    ocr_text = (
        "HD20260604 Date: 2026/4/22\n"
        "713041-2-V1-F1 Toy milk bottle 玩具牛奶瓶 528.32 KGM US$2.50 US$1,320.80"
    )
    monkeypatch.setattr(service.jakks_po_parser.pdfplumber, "open", lambda _path: FakePdf())
    monkeypatch.setattr(carton_mark, "get_rapidocr_engine", lambda: (lambda _image: object()))
    monkeypatch.setattr(carton_mark, "rapidocr_result_to_text", lambda _result: ocr_text)

    text, used_ocr = service.jakks_po_parser.extract_text("scan.pdf")

    assert text == ocr_text
    assert used_ocr is True


def test_jakks_inherits_unique_date_code_and_blocks_missing_chinese_name() -> None:
    order = {
        "lines": [{"item_no": "J-1", "product_name": "English Name"}],
        "warnings": [],
    }
    service._apply_jakks_schedule_data(
        order,
        {"records": [{"sku": "J-1", "product": "English Only", "date_code": "2628"}]},
    )

    line = order["lines"][0]
    assert line["date_code"] == "2628"
    assert any(flag["code"] == "missing_product_name_zh" for flag in line["flags"])
    assert not any(flag["code"] == "missing_date_code" for flag in line["flags"])


def test_jakks_native_contract_pdf_accepts_plain_item_and_trailing_usd(
    monkeypatch,
    tmp_path: Path,
) -> None:
    contract_text = """CONTRACT
JAKKS PACIFIC (H.K. ) LIMITED
Contract #: 811626
Printed Date: 11-Aug-2026 Order Date: 11-Aug-2026
Customer PO: 236856
Confirmation No: 67078 UNITED STATES
Ultimate Consignee Name: JAKKS PACIFIC, INC.
ITEM NUMBER DESCRIPTION QUANTITY PRICE EXTENDED PRICE
10022J CHARMING CUPID CUTIE PLUSH 5L 5L 100 3.0670 306.70 USD
Customer Item No.: 10022J 25 CA
UPC: 199460000214 Outer Pack: 4
Stock #: 10022J
Country of Origin: CHINA Line#: 1.000 Request Date: 13-Oct-2026
NOTES: OPEN WINDOW BOX
TOTAL USD 306.70
"""
    monkeypatch.setattr(
        service.jakks_po_parser,
        "extract_text",
        lambda _path: (contract_text, False),
    )

    parsed = service.jakks_po_parser.parse_po(tmp_path / "JDCUS-SI67078-OK811626-HE.pdf")

    assert parsed["contract_no"] == "811626"
    assert parsed["order_date"] == "2026-08-11"
    assert parsed["customer_po"] == "236856"
    assert parsed["confirmation_no"] == "67078"
    assert parsed["country"] == "UNITED STATES"
    assert parsed["customer"] == "JAKKS PACIFIC, INC. UNITED STATES"
    assert parsed["source_format"] == "pdf"
    assert parsed["used_ocr"] is False
    assert len(parsed["lines"]) == 1
    line = parsed["lines"][0]
    assert line == {
        "item_no": "10022J",
        "product_name": "CHARMING CUPID CUTIE PLUSH 5L 5L",
        "po_description": "CHARMING CUPID CUTIE PLUSH 5L 5L",
        "product_name_source": "PO英文品名",
        "quantity": 100,
        "unit": "PCS",
        "inner_pack": "",
        "outer_pack": 4,
        "cartons": 25,
        "unit_price_usd": 3.067,
        "total_usd": 306.7,
        "ship_date": "2026-10-13",
        "special_note": "OPEN WINDOW BOX",
        "product_packaging": "OPEN WINDOW BOX",
    }
    assert any("标准 CONTRACT PDF" in warning for warning in parsed["warnings"])


def test_jakks_export_overwrites_copied_po_field_formula(tmp_path: Path) -> None:
    schedule = _workbook_bytes(
        "26-Jakks排货表",
        {1: "来单日期", 5: "客户PO", 6: "合同号", 7: "客名", 9: "货号", 10: "产品名称", 11: "数量"},
        {5: "OLD", 6: "OLD", 7: "JAKKS", 9: "J-OLD", 10: "旧品名", 11: "=5000-4999"},
    )
    line = {
        "order_date": "2026-08-18", "customer_po": "236856", "contract_no": "811626",
        "customer": "JAKKS PACIFIC, INC. UNITED STATES", "item_no": "10022J",
        "product_name": "CHARMING CUPID CUTIE PLUSH 5L 5L", "quantity": 100,
        "unit": "PCS", "unit_price_usd": 3.067, "total_usd": 306.7,
        "source_file": "JDCUS-SI67078-OK811626-HE.pdf",
    }
    output = tmp_path / "jakks-formula-override.xlsx"

    service._export_prepared(
        customer_code="jakks",
        prepared=service.PreparedBatch(
            [line], [], "26-Jakks排货表总 Ai",
            export_payload={"filename": "Jakks批量", "lines": [line]},
        ),
        schedule_file_name="2026年Jakks排货表.xlsx",
        schedule_content=schedule,
        output_path=output,
    )

    workbook = openpyxl.load_workbook(output, data_only=False)
    try:
        worksheet = workbook["26-Jakks排货表"]
        assert worksheet.cell(3, 6).value == "811626"
        assert worksheet.cell(3, 9).value == "10022J"
        assert worksheet.cell(3, 11).value == 100
    finally:
        workbook.close()


def test_simba_prefers_same_name_excel_and_inherits_schedule(monkeypatch) -> None:
    parsed_names: list[str] = []
    monkeypatch.setattr(
        service.simba_schedule,
        "read_schedule",
        lambda *_args, **_kwargs: {"sheet": "Simba排期", "records": [{"contract_no": "S1", "item_no": "A1"}]},
    )

    def parse_po_file(_content, *, filename: str):
        parsed_names.append(filename)
        return {
            "filename": filename,
            "rows": [{
                "contract_no": "S1", "customer_po": "PO-S1", "item_no": "A1",
                "quantity": 120, "outer_pack": 12, "po_ship_date": "2026-09-01",
                "customer": "SIMBA", "contact": "May", "source_sheet": "转换PO",
            }],
            "warnings": [],
        }

    monkeypatch.setattr(service.simba_po_parser, "parse_po_file", parse_po_file)
    monkeypatch.setattr(
        service.simba_schedule,
        "enrich_rows_from_schedule",
        lambda rows, _schedule: {
            "product_names_applied": 1,
            "exact_fields_applied": 2,
            "product_name_conflicts": [],
        },
    )
    prepared = service._prepare_simba(
        [("Release-S1.pdf", b"pdf"), ("Release-S1.xlsx", b"xlsx")],
        "2026年Simba排期.xlsx",
        b"schedule",
        "2026-08-03",
    )
    assert parsed_names == ["Release-S1.xlsx"]
    assert len(prepared.records) == 1
    assert prepared.records[0]["order_date"] == "2026-08-03"
    assert any("优先采用 Excel" in warning for warning in prepared.warnings)


def test_simba_drops_po_carton_dimensions_and_keeps_first_english_sentence(monkeypatch) -> None:
    monkeypatch.setattr(
        service.simba_schedule,
        "read_schedule",
        lambda *_args, **_kwargs: {"sheet": "Simba排期", "records": []},
    )
    monkeypatch.setattr(
        service.simba_schedule,
        "enrich_rows_from_schedule",
        lambda *_args: {
            "product_names_applied": 0,
            "exact_fields_applied": 0,
            "product_name_conflicts": [],
        },
    )
    monkeypatch.setattr(
        service.simba_po_parser,
        "parse_po_file",
        lambda _content, *, filename: {
            "filename": filename,
            "rows": [{
                "contract_no": "S1",
                "item_no": "A1",
                "quantity": 24,
                "outer_pack": 12,
                "po_ship_date": "2026-09-01",
                "customer": "SIMBA",
                "contact": "May",
                "english_name": "Fire Truck. Includes two figures and accessories.",
                "outer_length_cm": 40,
                "outer_width_cm": 30,
                "outer_height_cm": 20,
            }],
            "warnings": [],
        },
    )

    prepared = service._prepare_simba(
        [("Release-S1.xlsx", b"xlsx")],
        "2026年Simba排期.xlsx",
        b"schedule",
        "2026-08-03",
    )

    row = prepared.records[0]
    assert row["english_name"] == "Fire Truck."
    assert "outer_length_cm" not in row
    assert "outer_width_cm" not in row
    assert "outer_height_cm" not in row


def test_simba_scanned_pdf_uses_shared_tesseract_configuration(monkeypatch) -> None:
    class FakePage:
        def extract_text(self, **_kwargs):
            return ""

        def to_image(self, **_kwargs):
            return type("Rendered", (), {"original": object()})()

    class FakePdf:
        pages = [FakePage()]

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

    configured: list[object] = []
    monkeypatch.setattr(common_pdf_text.pdfplumber, "open", lambda _source: FakePdf())
    monkeypatch.setattr(
        common_pdf_text,
        "configure_tesseract",
        lambda module: (configured.append(module) or (r"C:\Tesseract\tesseract.exe", "eng")),
    )
    monkeypatch.setattr(
        pytesseract,
        "image_to_string",
        lambda _image, *, lang, config: f"Release Order {lang} {config}",
    )

    text, pages, used_ocr = common_pdf_text.extract_pdf_text(b"scanned")

    assert configured == [pytesseract]
    assert text == "Release Order eng --psm 6"
    assert pages == 1
    assert used_ocr is True


def test_simba_release_parser_tolerates_ocr_reference_and_zero_packing() -> None:
    text = """
Release Order
Reference: $C700144914/ 800
Date of creation: 27.APR.2026
Mat. No.: 109491008
SPB Feature Plush, 30cm
Packing: OPC / 12PC Volume: 2.568 FT3 - 0.073 M3
Carton dimension: 59.500 x 26.000 x 47.000 CM
Quantity Master Contract No. PO Contract No. Unit Price Delivery Date
264 PC 500055145/ 10 300493420/ 10 40.45 HKD 09.JUN.2026
Total CTN: 22
Port of discharge: MELBOURNE
"""

    rows = simba_po_parser.parse_release_order_text(text, "RR_700144914.pdf")

    assert rows is not None
    assert len(rows) == 1
    assert rows[0]["contract_no"] == "700144914/800"
    assert rows[0]["item_no"] == "109491008"
    assert rows[0]["inner_pack"] == 0
    assert rows[0]["outer_pack"] == 12
    assert rows[0]["cartons"] == 22
    assert rows[0]["quantity"] == 264


def test_simba_empty_batch_reports_each_file_parser_failure(monkeypatch) -> None:
    monkeypatch.setattr(
        service.simba_schedule,
        "read_schedule",
        lambda *_args, **_kwargs: {"sheet": "Simba排期", "records": []},
    )
    monkeypatch.setattr(
        service.simba_po_parser,
        "parse_po_file",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(ValueError("OCR未配置")),
    )

    with pytest.raises(
        service.HuadengCustomerOrderError,
        match="RR700147106.pdf.*OCR未配置",
    ):
        service._prepare_simba(
            [("RR700147106.pdf", b"pdf")],
            "仙霸排货表.xlsx",
            b"schedule",
            "2026-08-12",
        )


def test_simba_export_adds_one_total_row_after_each_item_group(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "2026年排货表"
    headers = {
        2: "来单日期",
        5: "客户PO",
        6: "合同号",
        7: "客名",
        9: "货号",
        10: "产品名称",
        11: "数量",
        13: "外装箱",
        14: "总箱数",
    }
    for column, value in headers.items():
        worksheet.cell(1, column, value)
    for column, value in {
        2: datetime(2026, 7, 1),
        5: "PO-OLD",
        6: "OLD/100",
        7: "SIMBA",
        9: "ITEM-OLD",
        10: "旧产品",
        11: 10,
        13: 2,
        14: "=K2/M2",
    }.items():
        worksheet.cell(2, column, value)
    worksheet["J3"] = "合计："
    worksheet["K3"] = "=SUM(K2:K2)"
    worksheet["J3"].fill = openpyxl.styles.PatternFill("solid", fgColor="FFF2CC")
    worksheet["K3"].fill = openpyxl.styles.PatternFill("solid", fgColor="FFF2CC")
    source = BytesIO()
    workbook.save(source)
    workbook.close()
    output_path = tmp_path / "Simba逐款合计.xlsx"

    records = [
        {
            "order_date": "2026-08-14",
            "customer_po": "PO-A1",
            "contract_no": "A/100",
            "customer": "SIMBA",
            "item_no": "ITEM-A",
            "product_name": "产品A",
            "quantity": 12,
            "outer_pack": 3,
            "cartons": 4,
        },
        {
            "order_date": "2026-08-14",
            "customer_po": "PO-A2",
            "contract_no": "A/200",
            "customer": "SIMBA",
            "item_no": "ITEM-A",
            "product_name": "产品A",
            "quantity": 18,
            "outer_pack": 3,
            "cartons": 6,
        },
        {
            "order_date": "2026-08-14",
            "customer_po": "PO-B1",
            "contract_no": "B/100",
            "customer": "SIMBA",
            "item_no": "ITEM-B",
            "product_name": "产品B",
            "quantity": 20,
            "outer_pack": 4,
            "cartons": 5,
        },
    ]
    service._export_prepared(
        customer_code="simba",
        prepared=service.PreparedBatch(records, [], "2026年排货表"),
        schedule_file_name="Simba排期.xlsx",
        schedule_content=source.getvalue(),
        output_path=output_path,
    )

    rendered = openpyxl.load_workbook(output_path, data_only=False)
    worksheet = rendered["2026年排货表"]
    assert worksheet["K3"].value == "=SUM(K2:K2)"
    assert worksheet["I4"].value == worksheet["I5"].value == "ITEM-A"
    assert worksheet["J6"].value == "合计："
    assert worksheet["K6"].value == "=SUM(K4:K5)"
    assert worksheet["I7"].value == "ITEM-B"
    assert worksheet["J8"].value == "合计："
    assert worksheet["K8"].value == "=SUM(K7:K7)"
    assert worksheet["J6"].fill.fgColor.rgb == worksheet["J3"].fill.fgColor.rgb
    rendered.close()


def test_spin_master_uses_independent_schedule_and_composite_dedupe(monkeypatch) -> None:
    monkeypatch.setattr(
        service.spin_master_parser,
        "validate_schedule_file",
        lambda _path: {"sheet": "SPIN排期", "rows": [{"contract_no": "OLD"}]},
    )
    monkeypatch.setattr(
        service.spin_master_parser,
        "parse_po_file",
        lambda path: {"source": str(path), "items": [{"line_no": 10}]},
    )
    monkeypatch.setattr(
        service.spin_master_new_order_writer,
        "_records",
        lambda order: [{
            "customer_po": "PO-SM", "contract_no": "45001/10", "customer": "SPIN MASTER",
            "item_no": "GML/100/200", "product_name": "GML4pkSLD", "english_name": "GML4pkSLD",
            "quantity": 100, "outer_pack": 4, "cartons": 25, "ship_date": "2026-09-01",
            "unit_price_usd": 1, "total_usd": 100, "unit_price_hkd": 7.75, "total_hkd": 775,
            "source_file": order["source"],
        }],
    )
    prepared = service._prepare_spin_master(
        [("SM-1.pdf", b"one"), ("SM-1-copy.pdf", b"two")],
        "2026年SpinMaster排期.xlsx",
        b"schedule",
    )
    assert prepared.sheet_name == "SPIN排期"
    assert len(prepared.records) == 1
    assert any("完全重复" in warning for warning in prepared.warnings)


def test_spin_and_spin_master_specs_remain_separate() -> None:
    spin = service.get_huadeng_customer_mapping("spin")
    spin_master = service.get_huadeng_customer_mapping("spin-master")
    assert spin.target_template == "HUADENG_SPIN_SCHEDULE_APPEND_V2"
    assert spin_master.target_template == "HUADENG_SPIN_MASTER_SCHEDULE_APPEND_V2"
    assert spin.schedule_extensions == (".xlsx",)
    assert spin_master.schedule_extensions == (".xls", ".xlsx")


def test_spin_preview_uses_spin_schedule_family_and_keeps_manual_dates_blank(monkeypatch) -> None:
    schedule = _workbook_bytes(
        service.spin_schedule.MASTER_SHEET,
        {2: "来单期", 4: "联系人", 5: "客PO", 6: "合同号", 7: "客名", 8: "版本", 9: "货号", 10: "产品名称", 11: "数量", 13: "外箱", 27: "验货期", 28: "PO走货期", 36: "英文品名"},
        {5: "OLD", 6: "OLD/10", 7: "SPIN MASTER", 9: "1111/123456/7890/9999-10", 10: "Spin标准品名", 11: 100, 13: 4, 36: "Spin Product"},
    )

    def parse(_self, _path: str):
        return {
            "po_number": "45000001",
            "po_date": "2026-08-01",
            "contact": "Karen Sun",
            "customer": "SPIN MASTER",
            "lines": [{
                "line_no": "10", "qty": 100, "unit": "PC", "material_number": "123456",
                "description_en": "GML4pkSLD Spin Product", "sales_material": "7890",
                "material_group": "1111", "unit_price": 1.5, "total_usd": 150,
                "ship_date": "2026-09-01", "sales_order": "9999", "line_item": "10",
                "customer": "SPIN MASTER", "customer_po": "PO-SPIN",
                "item_key": "1111/123456/7890",
            }],
            "raw_text": "",
        }

    monkeypatch.setattr(service.spin_po_parser.SpinPOParser, "parse", parse)
    preview = service.create_huadeng_customer_preview(
        customer_code="spin",
        factory_id="huadeng",
        received_date="2026-08-03",
        po_files=[("Spin PO.pdf", b"pdf")],
        schedule_file_name="2026年Spin排期.xlsx",
        schedule_content=schedule,
    )
    row = preview["rows"][0]
    assert row["target_template"] == "HUADENG_SPIN_SCHEDULE_APPEND_V2"
    assert row["contract_no"] == "45000001/10"
    assert row["product_name_zh"] == "Spin标准品名"
    assert row["units_per_carton"] == "4"
    assert row["line_q"] == ""
    assert row["requested_ship_date"] == "2026-09-01"


def test_remaining_huadeng_exporters_append_to_complete_schedule_workbooks(tmp_path: Path) -> None:
    jakks_template = _workbook_bytes(
        "26-Jakks排货表总 Ai",
        {1: "来单日期", 5: "客户PO", 6: "合同号", 7: "客名", 9: "货号", 10: "产品名称", 11: "数量"},
        {5: "OLD", 6: "OLD", 7: "JAKKS", 9: "J-OLD", 10: "旧品名", 11: 1},
    )
    jakks_line = {
        "order_date": "2026-08-03", "customer_po": "PO-J1", "contract_no": "J1",
        "customer": "JAKKS", "item_no": "J-1", "product_name": "Jakks产品", "quantity": 20,
        "unit": "PCS", "unit_price_usd": 1, "total_usd": 20, "source_file": "J1.pdf",
    }
    jakks_output = tmp_path / "jakks-new.xlsx"
    service._export_prepared(
        customer_code="jakks",
        prepared=service.PreparedBatch(
            [jakks_line], [], "26-Jakks排货表总 Ai",
            export_payload={"filename": "Jakks批量", "lines": [jakks_line]},
        ),
        schedule_file_name="Jakks排期.xlsx",
        schedule_content=jakks_template,
        output_path=jakks_output,
    )

    simba_template = _workbook_bytes(
        "Simba排期",
        {2: "来单日期", 5: "客户PO", 6: "合同号", 7: "客名", 9: "货号", 10: "产品名称", 11: "数量"},
        {5: "OLD", 6: "OLD", 7: "SIMBA", 9: "S-OLD", 10: "旧品名", 11: 1},
    )
    simba_output = tmp_path / "simba-new.xlsx"
    service._export_prepared(
        customer_code="simba",
        prepared=service.PreparedBatch([{
            "order_date": "2026-08-03", "customer_po": "PO-S1", "contract_no": "S1",
            "customer": "SIMBA", "item_no": "S-1", "product_name": "Simba产品", "quantity": 24,
            "outer_pack": 12, "cartons": 2, "po_ship_date": "2026-09-01", "source_file": "S1.xlsx",
        }], [], "Simba排期"),
        schedule_file_name="Simba排期.xlsx",
        schedule_content=simba_template,
        output_path=simba_output,
    )

    spin_master_template = _workbook_bytes(
        "SPIN排期",
        {1: "来单期", 4: "客PO", 5: "合同号", 6: "客名", 8: "货号/SI", 9: "产品名称", 10: "数量"},
        {4: "OLD", 5: "OLD", 6: "SPIN MASTER", 8: "SM-OLD", 9: "旧品名", 10: 1},
    )
    spin_master_output = tmp_path / "spin-master-new.xlsx"
    service._export_prepared(
        customer_code="spin-master",
        prepared=service.PreparedBatch([{
            "order_date": "2026-08-03", "customer_po": "PO-SM", "contract_no": "45001/10",
            "customer": "SPIN MASTER", "item_no": "GML/100/200", "product_name": "Spin产品",
            "english_name": "Spin Product", "quantity": 100, "outer_pack": 4, "cartons": 25,
            "ship_date": "2026-09-01", "source_file": "SM.pdf",
        }], [], "SPIN排期"),
        schedule_file_name="SpinMaster排期.xlsx",
        schedule_content=spin_master_template,
        output_path=spin_master_output,
    )

    spin_template = _workbook_bytes(
        service.spin_schedule.MASTER_SHEET,
        {5: "客PO", 6: "合同号", 7: "客名", 9: "货号", 10: "产品名称", 11: "数量"},
        {5: "OLD", 6: "OLD", 7: "SPIN MASTER", 9: "SPIN-OLD", 10: "旧品名", 11: 1},
    )
    spin_output = tmp_path / "spin-new.xlsx"
    spin_values = {
        service.spin_schedule.COL["customer_po"]: "PO-SPIN",
        service.spin_schedule.COL["contract"]: "45001/10",
        service.spin_schedule.COL["customer"]: "SPIN MASTER",
        service.spin_schedule.COL["item"]: "100/200/300",
        service.spin_schedule.COL["cn_name"]: "Spin产品",
        service.spin_schedule.COL["qty"]: 100,
    }
    service._export_prepared(
        customer_code="spin",
        prepared=service.PreparedBatch(
            [], [], service.spin_schedule.MASTER_SHEET,
            legacy_rows=[spin_values], inheritance={},
        ),
        schedule_file_name="Spin排期.xlsx",
        schedule_content=spin_template,
        output_path=spin_output,
    )

    expected = (
        (jakks_output, "26-Jakks排货表总 Ai", 6, "J1"),
        (simba_output, "Simba排期", 6, "S1"),
        (spin_master_output, "SPIN排期", 5, "45001/10"),
        (spin_output, service.spin_schedule.MASTER_SHEET, 6, "45001/10"),
    )
    for output, sheet_name, contract_column, new_contract in expected:
        assert output.exists()
        workbook = openpyxl.load_workbook(output, data_only=False)
        try:
            assert workbook.sheetnames == [sheet_name]
            worksheet = workbook[sheet_name]
            assert worksheet.cell(2, contract_column).value == "OLD"
            assert worksheet.cell(3, contract_column).value == new_contract
        finally:
            workbook.close()
