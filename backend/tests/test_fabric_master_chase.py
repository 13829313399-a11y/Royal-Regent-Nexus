import importlib
from datetime import datetime
from io import BytesIO
from uuid import uuid4
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
from sqlalchemy import select
from test_fabric_procurement import BASE, book, row, stored
from test_fabric_procurement_tracking import save, undo_preview
from test_fabric_receiving import request, receive
from test_molding_sample_api import make_client, login_as

MASTER = "/api/fabric-warehouse/master"


def resolve(client, line, amount="37", **changes):
    payload = {"factory_id": "huakang-c", "request_id": str(uuid4()), "expected_source_revision": line["revision"],
        "expected_receipt_count": line["receipt_count"], "expected_resolution_revision": (line.get("chase_resolution") or {}).get("revision", 0),
        "starting_quantity": amount, "evidence": "采购核对原始欠数，包含系统已收10", "confirmed_start": True, **changes}
    return client.post(f"{BASE}/lines/{line['id']}/resolve-chase", json=payload), payload


def master(kind="MATERIAL", code="000123", name="白色布 58寸", **changes):
    return {"factory_id": "huakang-c", "request_id": str(uuid4()), "expected_revision": 0, "kind": kind, "code": code, "name": name,
            "status": "ACTIVE", "data": {"category": "FABRIC", "unit": "码"} if kind == "MATERIAL" else {}, **changes}


def test_unknown_physical_receipt_stays_unknown_until_audited_start_and_never_double_subtracts(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        batch = save(client, book([row(采购明细ID="K1", 订单数量=665, 入库数量="#VALUE!", 交货明细="")]))
        line = stored(client)["items"][0]
        detail = client.get(f"{BASE}/lines/{line['id']}", params={"factory_id": "huakang-c"}).json()
        assert detail["can_receive"] and detail["starting_chase_quantity"] is None
        posted = receive(client, line, request(line, batches=[{"quantity": "10", "location": "A1", "dye_lot": "0001"}]))
        assert posted.status_code == 200, posted.text
        assert posted.json()["prior_received_quantity"] is None
        save(client, book([row(采购明细ID="K1", 订单数量=665, 入库数量=10, 交货明细="")]))
        line = stored(client)["items"][0]
        assert line["warehouse_received_quantity"] == "10" and line["warehouse_outstanding_quantity"] is None
        response, payload = resolve(client, line)
        assert response.status_code == 200, response.text
        assert client.post(f"{BASE}/lines/{line['id']}/resolve-chase", json=payload).json() == response.json()
        assert client.post(f"{BASE}/lines/{line['id']}/resolve-chase", json={**payload, "evidence": "changed"}).status_code == 409
        assert resolve(client, line)[0].status_code == 409
        line = stored(client)["items"][0]
        assert line["starting_chase_quantity"] == "37" and line["warehouse_outstanding_quantity"] == "27"
        assert not line["receipt_quantity_review_required"]
        assert receive(client, line, request(line, delivery_reference="DN2", batches=[{"quantity": "27", "location": "A1", "dye_lot": "0002"}])).status_code == 200
        line = stored(client)["items"][0]
        assert line["warehouse_received_quantity"] == "37" and line["warehouse_outstanding_quantity"] == "0"
        assert undo_preview(client, batch).status_code == 409
        assert client.get(BASE + "/stock", params={"factory_id": "huakang-c"}).json()["total"] == 2
        save(client, book([row(采购明细ID="K1", 订单数量=665, 入库数量=37, 交货明细="")]))
        assert stored(client)["items"][0]["starting_chase_quantity"] == "37"
        assert stored(client, view="OUTSTANDING")["total"] == 0
        save(client, book([row(采购明细ID="K1", 订单数量=665, 入库数量=37, 基本单位="米")]))
        line = stored(client)["items"][0]
        assert resolve(client, line)[0].status_code == 409
        assert receive(client, line, request(line, delivery_reference="DN3")).status_code == 409


def test_confirmed_zero_closes_chase_without_creating_stock_and_blocks_withdrawal(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        batch = save(client, book([row(采购明细ID="K1", 入库数量=None, 交货明细="")]))
        line = stored(client)["items"][0]
        assert resolve(client, line, "101")[0].status_code == 422
        assert resolve(client, line, "0", confirmed_start=False)[0].status_code == 422
        assert resolve(client, line, "0")[0].status_code == 200
        assert stored(client, view="OUTSTANDING")["total"] == 0
        assert client.get(BASE + "/stock", params={"factory_id": "huakang-c"}).json()["total"] == 0
        assert undo_preview(client, batch).status_code == 409
        assert client.get(f"{BASE}/lines/{line['id']}", params={"factory_id": "huakang-c"}).json()["chase_resolution_history"][0]["starting_quantity"] == "0"


def test_master_candidates_require_confirmation_preserve_receipt_snapshots_and_enforce_versions(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        save(client, book([row(采购明细ID="K1", 入库数量=0, 交货明细="")]))
        candidates = client.get(MASTER + "/candidates", params={"factory_id": "huakang-c", "kind": "MATERIAL"}).json()["items"]
        assert candidates[0]["code"] == "000123" and not candidates[0]["data"].get("category")
        assert client.get(MASTER, params={"factory_id": "huakang-c", "kind": "MATERIAL"}).json()["total"] == 0
        assert client.post(MASTER, json=master(data={"unit": "码"})).status_code == 422
        body = master()
        created = client.post(MASTER, json=body)
        assert created.status_code == 200, created.text
        record = created.json()
        assert client.post(MASTER, json=body).json() == record
        assert client.post(MASTER, json={**body, "name": "different"}).status_code == 409
        assert client.post(MASTER, json=master()).status_code == 409
        location = client.post(MASTER, json=master("LOCATION", "A1", "布料货架一", data={"warehouse": "布料仓"})).json()
        assert location["status"] == "ACTIVE"
        line = stored(client)["items"][0]
        detail = client.get(f"{BASE}/lines/{line['id']}", params={"factory_id": "huakang-c"}).json()
        assert detail["material_category"] == "FABRIC" and detail["available_locations"][0]["code"] == "A1"
        posted = receive(client, line, request(line, batches=[{"quantity": "10", "location": "A1", "dye_lot": "0001"}])).json()
        assert posted["master_references"]["material"]["revision"] == 1
        update = {**body, "id": record["id"], "expected_revision": 1, "request_id": str(uuid4()), "status": "INACTIVE"}
        assert client.post(MASTER, json=update).status_code == 200
        assert client.post(MASTER, json={**update, "request_id": str(uuid4())}).status_code == 409
        current = stored(client)["items"][0]
        assert receive(client, current, request(current, delivery_reference="DN2")).status_code == 422
        history = client.get(BASE + "/receipts", params={"factory_id": "huakang-c"}).json()["items"][0]
        assert history["master_references"]["material"]["status"] == "ACTIVE"
        assert len(client.get(f"{MASTER}/{record['id']}/history", params={"factory_id": "huakang-c"}).json()["items"]) == 2
        assert client.post(MASTER, json=master("UNIT", "convert", "换算", data={"ratio": "0"})).status_code == 422
        assert client.post(MASTER, json=master("UNIT", "convert", "换算", data={"ratio": "1.2"})).status_code == 422
        assert client.get(MASTER, params={"factory_id": "huakang-d", "kind": "MATERIAL"}).status_code == 422
        assert client.post(MASTER, json=master(factory_id="huakang-d")).status_code == 422
        assert client.post(MASTER, json=master("UNIT", "码", "码", data={})).status_code == 200
        assert client.post(MASTER, json=master("UNIT", "U2", "码", data={})).status_code == 409
        assert client.post(MASTER, json=master("UNIT", "U3", "U3", data={})).status_code == 200


def promise_book(rows):
    wb = load_workbook(BytesIO(book(rows)))
    ws = wb["未回物料"]
    for column, name in ((16, "供应商复期"), (17, "二次复期")):
        ws.cell(2, column, name)
        for index, values in enumerate(rows, 3):
            ws.cell(index, column, values.get(name, ""))
    result = BytesIO(); wb.save(result); return result.getvalue()


def test_overdue_uses_latest_reply_and_all_filters_sort_before_pagination(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        tracking = importlib.import_module("app.services.fabric_procurement_tracking")
        monkeypatch.setattr(tracking, "business_now", lambda: datetime(2026, 10, 7, tzinfo=ZoneInfo("Asia/Shanghai")))
        rows = [row(采购明细ID=f"K{i}", 订单号=f"PO{i}", 入库数量=0, 交货明细="", 供应商复期=reply) for i, reply in enumerate(("2026-10-01", "2026-10-05", "2026-10-07", "2026-10-08", "等通知"))]
        rows.append(row(采购明细ID="K5", 订单号="PO5", 入库数量=100, 交货明细="", 供应商复期="2026-10-01"))
        rows[4]["二次复期"] = "等通知"; rows[4]["供应商复期"] = "2026-09-01"
        batch = save(client, promise_book(rows))
        result = stored(client, view="OUTSTANDING", sort="OVERDUE", limit=2, offset=1)
        assert result["total"] == 5 and [item["facts"]["order_no"] for item in result["items"]] == ["PO1", "PO2"]
        assert result["summary"]["overdue"] == 2 and result["summary"]["due_today"] == 1 and result["summary"]["awaiting_date"] == 1
        waiting = stored(client, view="AWAITING_DATE")["items"][0]
        assert waiting["promise_date"] is None
        assert stored(client, view="ARRIVAL_REVIEW")["items"][0]["overdue_days"] == 0
        assert stored(client, view="OUTSTANDING", promise_from="2026-10-05", promise_to="2026-10-07", import_batch=batch["id"], material_code="000123", supplier="采购供应商", source_category="PURCHASE", quantity_review="KNOWN")["total"] == 2
        assert stored(client, view="OUTSTANDING", category="UNCLASSIFIED")["total"] == 5
        assert client.get(BASE + "/lines", params={"factory_id": "huakang-c", "sort": "QUANTITY_DESC"}).status_code == 422
        assert stored(client, view="OUTSTANDING", sort="QUANTITY_DESC", unit="码")["total"] == 5
        assert client.get(BASE + "/lines", params={"factory_id": "huakang-c", "promise_from": "2026-10-07", "promise_to": "2026-10-01"}).status_code == 422
        first = stored(client, order_no="PO0")["items"][0]
        assert receive(client, first, request(first, batches=[{"quantity": "10", "location": "A1", "dye_lot": "0001"}])).status_code == 200
        assert stored(client, view="OVERDUE", order_no="PO0")["items"][0]["warehouse_outstanding_quantity"] == "90"


def test_master_and_chase_use_import_permission_and_explicit_denies(monkeypatch):
    from dataclasses import replace
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        save(client, book([row()]))
        line = stored(client)["items"][0]
        from test_raw_material_api import create_fixed_position_user
        create_fixed_position_user("fabric_manager", "position_warehouse_keeper", "pmc-warehouse", factory_id="huakang-c")
        auth = importlib.import_module("app.services.auth")
        dbm = importlib.import_module("app.db")
        models = importlib.import_module("app.models.auth")
        with dbm.SessionLocal() as db:
            actor = auth.build_auth_context(db, db.get(models.AuthUser, "user-fabric_manager"))
        denied = replace(actor, overrides=(auth.AuthOverrideContext(id="deny", permission_code="fabric_warehouse:import", effect="deny", factory_id="huakang-c", department="pmc-warehouse"),))
        client.app.dependency_overrides[auth.get_current_user] = lambda: denied
        assert client.get(MASTER, params={"factory_id": "huakang-c", "kind": "MATERIAL"}).json()["can_manage"] is False
        assert client.post(MASTER, json=master()).status_code == 403
        assert resolve(client, line)[0].status_code == 403
        wrong_factory = replace(actor, grants=tuple(replace(grant, factory_id="huakang-d") for grant in actor.grants))
        client.app.dependency_overrides[auth.get_current_user] = lambda: wrong_factory
        assert client.get(MASTER, params={"factory_id": "huakang-c", "kind": "MATERIAL"}).status_code == 403
        assert resolve(client, line)[0].status_code == 403
