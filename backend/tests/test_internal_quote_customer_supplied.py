import copy
import json
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.services.internal_quote import _rr2_cost_summary
from app.services.internal_quote_calculator import CalculationInputError, calculate_section
from app.services.internal_quote_excel import _build_summary_sheet
from test_internal_quote_excel_formula_links import _section, _find_row

SNAPSHOT = {"fx": {"rmb_hkd": ".85", "hkd_usd": "7.8", "rmb_usd": "7.75"}, "markup": "1.2"}


def sales_payload(rows=None, markup="1.2", misc=".03"):
    return {"pricing_mode": "component", "pricing_components": [{"id": "a", "name": "主体", "markup_x": markup}, {"id": "b", "name": "电话", "markup_x": "1.5"}],
            "shipping": {"markup_x": markup, "packaging_markup_x": "1.3", "misc_ratio": misc},
            "freight_calc": {"enabled": False},
            "customer_supplied_materials": rows if rows is not None else [{"item": "眼睛饰品", "unit_price_hkd": ".8", "fee_rate_percent": "3", "pricing_component_id": "a"}]}


def fixture(payload):
    engineering_payload = {"materials": [{"item": "主体五金", "category": "hardware", "quantity": 1, "unit_price_hkd": 10, "pricing_component_id": "a"},
                                         {"item": "电话五金", "category": "hardware", "quantity": 1, "unit_price_hkd": 5, "pricing_component_id": "b"}]}
    sales_calc = calculate_section("sales", payload, SNAPSHOT, "CUSTOMER-SUPPLIED")
    engineering_calc = calculate_section("engineering", engineering_payload, SNAPSHOT, "CUSTOMER-SUPPLIED")
    sections = [_section("sales", payload=payload, calculation=sales_calc), _section("engineering", payload=engineering_payload, calculation=engineering_calc)]
    for section in sections:
        section.is_required = True
    fee = Decimal(sales_calc["totals"]["customer_supplied_hkd"])
    context = {"factory_price_hkd": Decimal(15) + fee, "hardware_hkd": Decimal(15), "customer_supplied_hkd": fee}
    summary = _rr2_cost_summary(sections, context, SNAPSHOT, factory_id="huakang-b")
    return sections, context, summary


@pytest.mark.parametrize("markup,misc", [("1.2", ".03"), ("2.5", ".4"), (".5", "0")])
def test_only_custody_fee_is_added_without_markup_misc_or_tax(markup, misc):
    payload = sales_payload(markup=markup, misc=misc)
    sections, context, summary = fixture(payload)
    _, _, without = fixture(sales_payload([], markup, misc))
    values = lambda summary, group: {row["key"]: Decimal(row["value"]) for row in summary[group]}
    assert values(summary, "t1")["base_price"] - values(without, "t1")["base_price"] == Decimal(".0240")
    assert values(summary, "t2")["misc"] == values(without, "t2")["misc"]
    assert values(summary, "t2")["other_buy"] - values(without, "t2")["other_buy"] == Decimal(".0240")
    assert summary["totals"]["total_deduction_hkd"] == without["totals"]["total_deduction_hkd"]
    assert summary["totals"]["rmb_purchase_cost_hkd"] == without["totals"]["rmb_purchase_cost_hkd"]
    assert summary["shipping_pricing"]["pricing_groups"] == without["shipping_pricing"]["pricing_groups"]
    assert summary["shipping_pricing"]["global_pricing"] == without["shipping_pricing"]["global_pricing"]
    assert summary["shipping_pricing"]["customer_supplied_pricing"]["quoted_hkd"] == "0.0240"
    assert context["factory_price_hkd"] == Decimal("15.0240")
    calc = json.loads(sections[0].calculation_json)
    assert calc["totals"]["customer_supplied_usd"] == "0.0031"
    assert not any(e["kind"] == "customer_supplied_material" for e in summary["shipping_pricing"]["pricing_entries"])


@pytest.mark.parametrize("field,value", [("unit_price_hkd", -1), ("fee_rate_percent", -1), ("fee_rate_percent", 101), ("fee_rate_percent", "NaN"), ("unit_price_hkd", "Infinity"), ("item", ""), ("unit_price_hkd", ""), ("fee_rate_percent", "")])
def test_customer_material_input_validation(field, value):
    payload = sales_payload()
    payload["customer_supplied_materials"][0][field] = value
    with pytest.raises(CalculationInputError):
        calculate_section("sales", payload, SNAPSHOT, "TEST")


def test_zero_and_multiple_materials_and_ordinary_compatibility():
    payload = sales_payload([{"item": "免费保管", "unit_price_hkd": 12, "fee_rate_percent": 0},
                             {"item": "饰品", "unit_price_hkd": .8, "fee_rate_percent": 3},
                             {"item": "附件", "unit_price_hkd": 2, "fee_rate_percent": 5}])
    result = calculate_section("sales", payload, SNAPSHOT, "TEST")
    assert result["totals"]["customer_supplied_hkd"] == "0.1240"
    assert len(result["line_breakdown"]) == 3
    payload.pop("pricing_mode")
    with pytest.raises(CalculationInputError, match="仅适用"):
        calculate_section("sales", payload, SNAPSHOT, "TEST")
    payload["customer_supplied_materials"] = []
    assert calculate_section("sales", payload, SNAPSHOT, "TEST")["totals"]["customer_supplied_hkd"] == "0.0000"


def test_small_custody_fees_use_the_same_half_up_rounding_as_excel():
    payload = sales_payload([{"item": "饰品", "unit_price_hkd": ".35", "fee_rate_percent": ".1"},
                             {"item": "饰片", "unit_price_hkd": ".15", "fee_rate_percent": ".1"}])
    result = calculate_section("sales", payload, SNAPSHOT, "TEST")
    assert [line["amount_hkd"] for line in result["line_breakdown"]] == ["0.0004", "0.0002"]
    assert result["totals"]["customer_supplied_hkd"] == "0.0006"


def test_excel_has_component_scoped_formula_blocks_no_multiplier_and_one_final_reference():
    payload = sales_payload([{"item": "眼睛饰品", "unit_price_hkd": ".8", "fee_rate_percent": 3, "pricing_component_id": "a"},
                             {"item": "电话饰片", "unit_price_hkd": 2, "fee_rate_percent": 5, "pricing_component_id": "b"}])
    sections, context, summary = fixture(payload)
    workbook = Workbook()
    _build_summary_sheet(workbook, SimpleNamespace(product_name="客供物料验证", quote_no="JP-CUSTODY", customer="JustPlay", region_code="mainland", remark="测试"), sections, SNAPSHOT, summary, context, [])
    sheet = workbook["报价明细"]
    main_start = _find_row(sheet, 3, "主体 · 客供物料")
    phone_start = _find_row(sheet, 3, "电话 · 客供物料")
    assert main_start < _find_row(sheet, 1, "电话-明细") < phone_start
    for start, formula in [(main_start, "=ROUND(0.8*0.03,4)"), (phone_start, "=ROUND(2*0.05,4)")]:
        assert sheet.cell(start + 1, 4).value == formula
        assert sheet.cell(start + 1, 1).value in (None, "")  # No refundable purchase tax tag.
        assert sheet.cell(start + 1, 2).value == "其他外购"
        assert sheet.cell(start + 2, 4).value == f"=SUM(D{start+1}:D{start+1})"
        assert sheet.cell(start + 3, 4).value == "=$R$4"
        assert sheet.cell(start + 4, 4).value == f"=D{start+2}/D{start+3}"
        assert not any(sheet.cell(row, 3).value in {"×", "÷"} for row in range(start, start + 5))
        assert sheet.cell(start, 1).border.top.style == "medium"
    total = _find_row(sheet, 3, "产品价（含运费）（USD）：")
    for start in (main_start, phone_start):
        assert sheet.cell(total, 4).value.count(f"$D${start+4}") == 1
    misc_row = _find_row(sheet, 3, "旺季价") + 4
    assert f"$D${main_start+2}" in sheet.cell(misc_row, 13).value
    assert f"$D${phone_start+2}" in sheet.cell(misc_row, 13).value
    buffer = BytesIO(); workbook.save(buffer)
    with_saved_formulas = load_workbook(BytesIO(buffer.getvalue()))
    assert with_saved_formulas["报价明细"].cell(main_start + 1, 4).value == "=ROUND(0.8*0.03,4)"
    with_saved_formulas.close(); workbook.close()


def test_api_save_preview_delete_and_history_component_reuse(monkeypatch):
    from test_internal_quote_api import make_client, login, create_payload, logout
    from test_internal_quote_history import catalog, source_ref, parts
    with make_client(monkeypatch) as client:
        login(client, "custody-sales", "sales_customer_owner", "sales-business", "huakang-b")
        body = create_payload(suffix="CUSTODY")
        body.update(factory_id="huakang-b", customer="JustPlay", pricing_components=["主体", "电话"])
        response = client.post("/api/internal-quotes", json=body); assert response.status_code == 201, response.text
        created = response.json(); qid = created["id"]
        payload = parts(created, "sales")
        payload["customer_supplied_materials"] = [{"item": "主体眼饰", "unit_price_hkd": .8, "fee_rate_percent": 3, "pricing_component_id": "component-01"},
                                                  {"item": "电话饰片", "unit_price_hkd": 2, "fee_rate_percent": 5, "pricing_component_id": "component-02"}]
        response = client.put(f"/api/internal-quotes/{qid}/sections/sales", json={"revision": 1, "payload": payload})
        assert response.status_code == 200, response.text
        assert response.json()["calculation"]["totals"]["factory_price_hkd"] == "0.1240"
        revision = response.json()["revision"]
        for _ in range(2):
            preview = client.post(f"/api/internal-quotes/{qid}/cost-preview", json={"drafts":[{"section_code":"sales","revision":revision,"payload":payload}]})
            assert preview.status_code == 200, preview.text
            assert preview.json()["preview_factory_price_hkd"] == "0.1240"
        old = next(r for r in catalog(client) if r["quote_id"] == qid)
        copied_body = copy.deepcopy(body)
        copied_body.update(quote_no="IQ-CUSTODY-COPY", products=[{"product_name":"电话改作主体", "qty":500,
            "pricing_components":["新主体"], "component_sources":[source_ref(old,"component-02")]}])
        copied = client.post("/api/internal-quotes", json=copied_body)
        assert copied.status_code == 201, copied.text
        assert parts(copied.json(), "sales")["customer_supplied_materials"] == [{"item":"电话饰片", "unit_price_hkd":2, "fee_rate_percent":5, "pricing_component_id":"component-01"}]
        payload["customer_supplied_materials"] = []
        removed = client.put(f"/api/internal-quotes/{qid}/sections/sales", json={"revision":revision,"payload":payload})
        assert removed.status_code == 200, removed.text
        assert removed.json()["calculation"]["totals"]["factory_price_hkd"] == "0.0000"
        logout(client)
        login(client,"custody-ordinary","sales_customer_owner","sales-business")
        ordinary = client.post("/api/internal-quotes",json=create_payload(suffix="CUSTODY-ORDINARY")).json()
        refused = client.put(f"/api/internal-quotes/{ordinary['id']}/sections/sales",json={"revision":1,"payload":sales_payload()})
        assert refused.status_code == 400 and "仅适用" in refused.text
