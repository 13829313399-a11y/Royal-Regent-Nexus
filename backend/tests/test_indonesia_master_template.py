"""Compatibility with the maintained 57-column shipment workbook, using synthetic data."""
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image
import pytest

from app.services.indonesia_master_import import parse_material_book
from app.services.indonesia_material_excel import TEMPLATE
from test_indonesia_materials import client, command, material, product, workspace
from test_indonesia_master_import import preview, choices

S = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def rewrite(content, change):
    with ZipFile(BytesIO(content)) as original:
        parts = {name: original.read(name) for name in original.namelist()}
    change(parts)
    output = BytesIO()
    with ZipFile(output, 'w', ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)
    return output.getvalue()


def template_book(duplicate_image=False):
    reference = openpyxl.load_workbook(TEMPLATE, read_only=True)
    headers = [cell.value for cell in next(reference['5.21出货明细'].iter_rows(min_row=3, max_row=3, max_col=57))]
    reference.close()
    book = openpyxl.Workbook(); sheet = book.active; sheet.title = '柜号可变'
    sheet.append(['出货资料']); sheet.append(headers)
    for n, name in enumerate(['外部货号物料', '子弹筒', '绿色拉杆', '蓝色拉杆', '推杆'], 3):
        for col, value in {'D': '99999' if n == 3 else '40149', 'F': name, 'G': 'Plastic part',
                           'I': '塑胶', 'J': 'KGM', 'M': 'pcs', 'L': 6844, 'AO': 4.66,
                           'AZ': '53003116', 'BD': '=1/3'}.items():
            sheet[f'{col}{n}'] = value
    for col, value in [('AL', '合并供应商'), ('AR', '合并出口公司')]:
        sheet[f'{col}3'] = value; sheet.merge_cells(f'{col}3:{col}7')
    sheet['D8'] = '40149'; sheet['F8'] = '区域外物料'; sheet['AZ8'] = 'OUTSIDE'
    # Independent photos are distinguishable even when their business codes coincide.
    for n in range(4, 8):
        photo = BytesIO(); Image.new('RGB', (16, 12), (n * 30, 50, 20)).save(photo, 'PNG')
        sheet.add_image(ExcelImage(photo), f'T{n}')
    if duplicate_image:
        photo = BytesIO(); Image.new('RGB', (16, 12), 'red').save(photo, 'PNG')
        sheet.add_image(ExcelImage(photo), 'T4')
    output = BytesIO(); book.save(output)

    def cache(parts):
        xml = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        for node in xml.findall('.//'+S+'c'):
            if node.get('r', '').startswith('BD') and node.find(S+'f') is not None:
                node.find(S+'v').text = '0.3333333333333333'
        parts['xl/worksheets/sheet1.xml'] = ET.tostring(xml)
    return rewrite(output.getvalue(), cache)


def test_template_units_actual_merged_ranges_cached_weights_and_row_photos():
    parsed = parse_material_book(template_book(), '任意文件.xlsx')
    assert parsed['sheets'][0]['header_row'] == 2
    assert parsed['sheets'][0]['columns']['customs_unit'] == 10
    assert parsed['sheets'][0]['columns']['unit'] == 13
    rows = {row['source_row']: row for row in parsed['rows']}
    for row in range(4, 8):
        assert rows[row]['source']['customs_unit'] == 'KGM'
        assert rows[row]['source']['unit'] == 'pcs'
        assert rows[row]['source']['supplier'] == '合并供应商'
        assert rows[row]['source']['export_company'] == '合并出口公司'
        assert rows[row]['source']['net_weight_kg'] == '0.33333333'
        assert rows[row]['evidence']['supplier'] == {'cell': 'AL3', 'merged': True}
        assert rows[row]['evidence']['net_weight_kg']['cached_value'] == '0.3333333333333333'
        assert 'piece_usage' not in rows[row]['source'] and 'currency' not in rows[row]['source']
    assert len({rows[row]['source']['image_data'] for row in range(4, 8)}) == 4
    assert 'supplier' not in rows[8]['source'] and 'export_company' not in rows[8]['source']
    assert 'image_data' not in rows[8]['source']


@pytest.mark.parametrize('cached,kind,expected', [('0','n','0.00000000'), ('0.123456785','n','0.12345679'),
                                               (None,'n',None), ('#DIV/0!','e',None), ('NaN','str',None), ('123','str',None)])
def test_formula_cache_zero_rounding_missing_errors_and_text(cached, kind, expected):
    def change(parts):
        root = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        cell = next(c for c in root.findall('.//'+S+'c') if c.get('r') == 'BD4')
        cell.set('t', kind); cell.find(S+'v').text = cached
        parts['xl/worksheets/sheet1.xml'] = ET.tostring(root)
    row = next(r for r in parse_material_book(rewrite(template_book(), change), '模板.xlsx')['rows'] if r['source_row'] == 4)
    assert row['source'].get('net_weight_kg') == expected
    assert bool(row['warnings']) == (expected is None)


@pytest.mark.parametrize('mode', ['multiple', 'external', 'invalid', 'wrong-column', 'cross-row'])
def test_ambiguous_external_and_invalid_pictures_do_not_drop_materials(mode):
    def change(parts):
        ns = {'x': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'}
        if mode == 'external':
            root = ET.fromstring(parts['xl/drawings/_rels/drawing1.xml.rels'])
            root[0].set('TargetMode', 'External'); root[0].set('Target', 'https://example.invalid/private.png')
            parts['xl/drawings/_rels/drawing1.xml.rels'] = ET.tostring(root)
        elif mode == 'invalid':
            parts['xl/media/image1.png'] = b'not a picture'
        elif mode in {'wrong-column', 'cross-row'}:
            root = ET.fromstring(parts['xl/drawings/drawing1.xml'])
            if mode == 'wrong-column':
                root[0].find('x:from/x:col', ns).text = '18'
            else:
                root[0].tag = '{'+ns['x']+'}twoCellAnchor'
                end = ET.SubElement(root[0], '{'+ns['x']+'}to')
                ET.SubElement(end, '{'+ns['x']+'}row').text = '5'
            parts['xl/drawings/drawing1.xml'] = ET.tostring(root)
    rows = parse_material_book(rewrite(template_book(mode == 'multiple'), change), '模板.xlsx')['rows']
    row = next(r for r in rows if r['source_row'] == 4)
    assert row['source']['name'] == '子弹筒' and 'image_data' not in row['source']
    assert len(rows) == 6
    assert all(r['source'].get('image_data') for r in rows if r['source_row'] in [5, 6, 7])
    if mode != 'wrong-column':
        assert row['warnings']


def test_template_import_preserves_keep_image_and_saves_new_photos(client):
    existing = material(client, code='53003116', name='子弹筒', specification='')
    p = product(client, existing, code='40149')
    staged = preview(client, template_book(), p)
    payload = choices(staged)
    payload['choices'] = payload['choices'][:4]
    for choice in payload['choices']:
        choice['material']['currency'] = 'USD'  # explicit test business input, absent from template
        choice['recipes'][0]['usage'] = '2'
    payload['choices'][0].update(decision='keep', material=None)
    before = deepcopy(existing)
    saved = command(client, 'master_import_confirm', payload, staged)
    records = {r['id']:r for r in workspace(client)['records']}
    assert records[existing['id']] == before
    for item in saved['data']['choices'][1:]:
        result = records[item['material_id']]['data']
        assert result['image_data'].startswith('data:image/png;base64,')
        assert result['supplier'] == '合并供应商'
        assert result['net_weight_kg'] == '0.33333333'


def test_explicit_mapping_can_omit_template_fields():
    from app.schemas.indonesia_materials import MasterSheetMapping
    mapping = MasterSheetMapping(sheet='柜号可变', header_row=2, columns={'item_no':4,'name':6,'code':52})
    parsed = parse_material_book(template_book(), '模板.xlsx', [mapping])
    assert 'unit' not in parsed['rows'][0]['source']
    assert 'image_data' not in parsed['rows'][1]['source']
