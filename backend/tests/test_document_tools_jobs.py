"""Real private uploads, durable leases, revisions and file delivery."""
import io
import time
from types import SimpleNamespace
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event, select, update
from sqlalchemy.orm import sessionmaker

from app.api.document_tools import router
from app.core.config import settings
from app.db import Base, get_db
from app.models.auth import AuthUser
from app.models.document_tools import DocumentToolSource as Source, DocumentToolJob as Job, DocumentToolArtifact as Artifact, DocumentToolCorrection as Correction
from app.services.auth import get_current_user
from app.services.document_tools import job_service as jobs, pipeline, storage


def test_simultaneous_password_initialization_uses_one_complete_key(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "document_tools_storage_dir", str(tmp_path / "passwords"))
    barrier = Barrier(16)
    def encrypt(index):
        barrier.wait()
        return storage.encrypt_password(f"private-password-{index}")
    with ThreadPoolExecutor(max_workers=16) as executor:
        encrypted = list(executor.map(encrypt, range(16)))
    assert [storage.decrypt_password(value) for value in encrypted] == [f"private-password-{i}" for i in range(16)]
    assert len(list((tmp_path / "passwords").iterdir())) == 1


@pytest.fixture
def runtime(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "document_tools_storage_dir", str(tmp_path / "files"))
    monkeypatch.setattr(settings, "document_tools_ai_mode", "off")
    # These contract tests exercise native extraction; true Office rendering has its own suite.
    monkeypatch.setattr(settings, "document_tools_office_command", "missing-office-contract-test")
    monkeypatch.delenv("DOCUMENT_TOOLS_SOFFICE_PATH", raising=False)
    monkeypatch.delenv("LIBREOFFICE_PATH", raising=False)
    engine = create_engine("sqlite:///" + str(tmp_path / "jobs.db"), connect_args={"check_same_thread": False})
    @event.listens_for(engine, "connect")
    def foreign_keys(connection, record):
        connection.execute("PRAGMA foreign_keys=ON")
    Base.metadata.create_all(engine, tables=[AuthUser.__table__, Source.__table__, Job.__table__, Artifact.__table__, Correction.__table__])
    sessions = sessionmaker(engine, expire_on_commit=False)
    with sessions() as db:
        db.add_all([AuthUser(id=x, username=x, password_salt="test", password_hash="test") for x in ("alice", "bob")])
        db.commit()
    app = FastAPI()
    app.include_router(router)
    current = SimpleNamespace(id="alice")
    def database():
        with sessions() as db:
            yield db
    app.dependency_overrides[get_db] = database
    app.dependency_overrides[get_current_user] = lambda: current
    with TestClient(app) as client:
        yield client, sessions, current
    engine.dispose()


def upload(client, name="table.docx", data=None):
    if data is None:
        doc = Document()
        doc.add_paragraph("Source notes 12.50")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Item"
        table.cell(0, 1).text = "Quantity"
        table.cell(1, 0).text = "000123"
        table.cell(1, 1).text = "12.50"
        buffer = io.BytesIO()
        doc.save(buffer)
        data = buffer.getvalue()
    response = client.post("/api/tools/uploads", files={"file": (name, data)}, data={"factory_id": "huakang-a"})
    assert response.status_code == 202, response.text
    return response.json()


def converted(runtime):
    client, sessions, _ = runtime
    source = upload(client)
    assert pipeline.run_one(sessions)
    inspected = client.get('/api/tools/jobs/' + source['inspection_job_id']).json()
    assert inspected["execution_status"] == "succeeded", inspected
    response = client.post('/api/tools/jobs', json={"source_id": source["source_id"], "operation": "word_to_excel", "options": {}, "client_request_id": jobs.uid()})
    assert response.status_code == 202, response.text
    job_id = response.json()["job_id"]
    pipeline.run_one(sessions)
    job = client.get('/api/tools/jobs/' + job_id).json()
    assert job["execution_status"] == "succeeded", job
    return source, job


def test_native_upload_convert_range_and_immutable_revision(runtime):
    client, sessions, _ = runtime
    source, job = converted(runtime)
    result = client.get(f'/api/tools/jobs/{job["id"]}/result').json()
    cell = next(c for t in result["tables"] for c in t["cells"] if c["display_text"] == "000123")
    artifact = next(a for a in job["artifacts"] if a["role"] == "result")
    original = client.get(f'/api/tools/artifacts/{artifact["id"]}/download').content
    partial = client.get(f'/api/tools/artifacts/{artifact["id"]}/content', headers={"Range": "bytes=0-9"})
    assert partial.status_code == 206 and partial.content == original[:10]
    revision = client.post(f'/api/tools/jobs/{job["id"]}/revise', json={"base_revision": 1, "corrections": [{"target_id": cell["id"], "new_value": "000456", "reason": "verified source"}]})
    assert revision.status_code == 202, revision.text
    pipeline.run_one(sessions)
    child = client.get('/api/tools/jobs/' + revision.json()["job_id"]).json()
    assert child["execution_status"] == "succeeded", child
    revised = client.get(f'/api/tools/jobs/{child["id"]}/result').json()
    assert any(c["display_text"] == "000456" for t in revised["tables"] for c in t["cells"])
    assert client.get(f'/api/tools/artifacts/{artifact["id"]}/download').content == original
    assert child["revision"] == 2
    with sessions() as db:
        correction = db.scalar(select(Correction))
        assert correction.old_value == "000123" and correction.author_user_id == "alice"


def test_owner_isolation_all_routes(runtime):
    client, _, current = runtime
    source, job = converted(runtime)
    artifact = job["artifacts"][0]
    current.id = "bob"
    for path in [f'sources/{source["source_id"]}', f'jobs/{job["id"]}', f'jobs/{job["id"]}/result', f'jobs/{job["id"]}/issues', f'artifacts/{artifact["id"]}/content', f'artifacts/{artifact["id"]}/download']:
        assert client.get('/api/tools/' + path).status_code == 404, path
    for action in ["cancel", "retry"]:
        assert client.post(f'/api/tools/jobs/{job["id"]}/{action}').status_code == 404
    assert client.delete(f'/api/tools/jobs/{job["id"]}').status_code == 404
    assert client.post('/api/tools/packages', json={"artifact_ids": [artifact["id"]], "client_request_id": jobs.uid()}).status_code == 404
    assert client.get('/api/tools/jobs').json()["total"] == 0


def test_idempotency_and_expired_lease_recovery(runtime):
    client, sessions, _ = runtime
    source = upload(client)
    first = jobs.claim_job(sessions)
    assert first and jobs.claim_job(sessions) is None
    with sessions() as db:
        db.execute(update(Job).where(Job.id == first[0]).values(lease_until=time.time() - 1))
        db.commit()
    recovered = jobs.claim_job(sessions)
    assert recovered[0] == first[0] and recovered[1] != first[1]
    assert not jobs.renew(sessions, *first)
    pipeline.run_one(sessions, recovered)
    payload = {"source_id": source["source_id"], "operation": "word_to_excel", "options": {}, "client_request_id": "one-submit"}
    a = client.post('/api/tools/jobs', json=payload)
    b = client.post('/api/tools/jobs', json=payload)
    assert a.status_code == b.status_code == 202 and a.json() == b.json()
    payload["options"] = {"word_mode": "structure"}
    assert client.post('/api/tools/jobs', json=payload).status_code == 409


def test_cancel_queued_retry_and_invalid_file(runtime):
    client, sessions, _ = runtime
    source = upload(client, data=b"fake zip")
    cancelled = client.post(f'/api/tools/jobs/{source["inspection_job_id"]}/cancel')
    assert cancelled.json()["execution_status"] == "cancelled"
    assert pipeline.run_one(sessions) is False
    retry = client.post(f'/api/tools/jobs/{source["inspection_job_id"]}/retry').json()["job_id"]
    pipeline.run_one(sessions)
    failure = client.get('/api/tools/jobs/' + retry).json()
    assert failure["execution_status"] == "failed" and failure["error_code"] == "INVALID_FILE"


def test_password_and_package(runtime):
    from pypdf import PdfWriter
    client, sessions, _ = runtime
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=300)
    writer.encrypt("sample-pass")
    buffer = io.BytesIO()
    writer.write(buffer)
    source = upload(client, "protected.pdf", buffer.getvalue())
    pipeline.run_one(sessions)
    assert client.get('/api/tools/sources/' + source["source_id"]).json()["inspection_status"] == "awaiting_input"
    assert client.post(f'/api/tools/sources/{source["source_id"]}/password', json={"password": "sample-pass"}).status_code == 202
    pipeline.run_one(sessions)
    assert client.get('/api/tools/sources/' + source["source_id"]).json()["inspection_status"] == "succeeded"
    with sessions() as db:
        assert "sample-pass" not in db.get(Source, source["source_id"]).credential_ciphertext
    _, job = converted(runtime)
    result = next(a for a in job["artifacts"] if a["role"] == "result")
    package = client.post('/api/tools/packages', json={"artifact_ids": [result["id"]], "client_request_id": jobs.uid()}).json()["job_id"]
    pipeline.run_one(sessions)
    assert client.get('/api/tools/jobs/' + package).json()["execution_status"] == "succeeded"


def test_path_guard(runtime):
    from app.services.document_tools.document_ir import ToolError
    with pytest.raises(ToolError):
        storage.resolve('../escape')


def test_manual_numeric_correction_preserves_excel_type(runtime, tmp_path):
    from openpyxl import Workbook, load_workbook
    from app.services.document_tools import office_engine
    source = tmp_path / 'typed.xlsx'
    workbook = Workbook()
    workbook.active.append(['Item', 'Amount'])
    workbook.active.append(['00123', 12.5])
    workbook.active['B2'].number_format = '0.00'
    workbook.save(source)
    ir = office_engine.inspect_office(source, {}, tmp_path / 'inspection', lambda *_: None, lambda: False).ir
    cells = {cell.source.cell: cell for table in ir.tables for cell in table.cells}
    pipeline.apply_corrections(ir, [SimpleNamespace(target_anchor={'target_id': cells[anchor].id}, new_value=value)
                                   for anchor, value in [('A2', '00456'), ('B2', '13.75')]])
    output = tmp_path / 'corrected.xlsx'
    office_engine.write_xlsx(ir, output, {})
    sheet = load_workbook(output).worksheets[0]
    assert sheet['A2'].value == '00456' and sheet['A2'].data_type == 's'
    assert sheet['B2'].value == 13.75 and sheet['B2'].data_type == 'n'
    assert sheet['B2'].number_format == '0.00'


def test_result_window_does_not_repeat_full_rendering_cache(runtime):
    client, sessions, _ = runtime
    _, job = converted(runtime)
    with sessions() as db:
        artifact = db.scalar(select(Artifact).where(Artifact.job_id == job["id"], Artifact.role == "ir"))
        key = artifact.storage_key
    ir = storage.read_json(key)
    ir['engine_manifest']['display_cells'] = {'Sheet1': {str(i): {'formatted': 'rendered value'} for i in range(10000)}}
    storage.atomic_json(storage.resolve(key), ir)
    response = client.get('/api/tools/jobs/' + job['id'] + '/result', params={'limit': 2})
    assert response.status_code == 200
    result = response.json()
    assert sum(len(table['cells']) for table in result['tables']) == 2
    assert 'display_cells' not in result['engine_manifest']
    assert result['engine_manifest']['display_cell_count'] == 10000
    assert len(response.content) < 20000


def test_stale_attempt_cannot_publish_failure_over_new_worker(runtime, monkeypatch):
    from app.services.document_tools import office_engine
    client, sessions, _ = runtime
    source = upload(client)
    def replaced_attempt(*args, **kwargs):
        with sessions() as db:
            db.execute(update(Job).where(Job.id == source["inspection_job_id"]).values(lease_token="new-worker", lease_until=time.time() + 90))
            db.commit()
        raise RuntimeError("old attempt crashed")
    monkeypatch.setattr(office_engine, "inspect_office", replaced_attempt)
    pipeline.run_one(sessions)
    with sessions() as db:
        job = db.get(Job, source["inspection_job_id"])
        assert job.execution_status == "running" and job.lease_token == "new-worker"


def test_expired_result_and_terminal_inspection_recovery(runtime):
    client, sessions, _ = runtime
    source, job = converted(runtime)
    with sessions() as db:
        db.execute(update(Artifact).where(Artifact.job_id == job["id"], Artifact.role == "ir").values(expires_at=time.time() - 1))
        db.commit()
    assert client.get(f'/api/tools/jobs/{job["id"]}/result').status_code == 410
    assert client.get(f'/api/tools/jobs/{job["id"]}/issues').status_code == 410
    new_source = upload(client)
    claimed = jobs.claim_job(sessions)
    with sessions() as db:
        db.execute(update(Job).where(Job.id == claimed[0]).values(lease_until=time.time() - 1, attempt=settings.document_tools_max_attempts))
        db.commit()
    assert jobs.claim_job(sessions) is None
    assert client.get('/api/tools/sources/' + new_source["source_id"]).json()["inspection_status"] == "failed"


def test_cancel_running_engine_preserves_other_jobs(runtime, monkeypatch):
    from app.services.document_tools import office_engine
    from app.services.document_tools.document_ir import Cancelled
    client, sessions, _ = runtime
    source = upload(client)
    def interrupted(*args, **kwargs):
        client.post(f'/api/tools/jobs/{source["inspection_job_id"]}/cancel')
        raise Cancelled()
    monkeypatch.setattr(office_engine, "inspect_office", interrupted)
    other = upload(client)
    pipeline.run_one(sessions)
    assert client.get('/api/tools/jobs/' + source["inspection_job_id"]).json()["execution_status"] == "cancelled"
    assert client.get('/api/tools/jobs/' + other["inspection_job_id"]).json()["execution_status"] == "queued"


@pytest.mark.parametrize("state", ["queued", "running", "awaiting_input"])
def test_withdraw_fences_offline_worker_and_allows_retry(runtime, state):
    client, sessions, _ = runtime
    source = upload(client)
    job_id = source["inspection_job_id"]
    with sessions() as db:
        db.execute(update(Job).where(Job.id == job_id).values(
            execution_status=state, lease_token="old-worker", lease_until=time.time() + 600))
        db.commit()
    response = client.post(f'/api/tools/jobs/{job_id}/cancel')
    assert response.status_code == 200
    assert response.json()["execution_status"] == "cancelled"
    assert not jobs.renew(sessions, job_id, "old-worker")
    with sessions() as db:
        assert db.scalar(select(Job).where(jobs.lease_filter(job_id, "old-worker"))) is None
    assert client.get(f'/api/tools/sources/{source["source_id"]}').json()["inspection_status"] == "cancelled"
    assert client.post(f'/api/tools/jobs/{job_id}/cancel').json()["execution_status"] == "cancelled"
    assert client.post(f'/api/tools/jobs/{job_id}/retry').status_code == 202
    assert pipeline.run_one(sessions)


@pytest.mark.parametrize("state", ["queued", "running", "awaiting_input", "failed", "cancelled", "succeeded"])
def test_delete_is_durable_private_idempotent_and_preserves_source(runtime, state):
    client, sessions, current = runtime
    source = upload(client)
    other = upload(client)
    job_id = source["inspection_job_id"]
    with sessions() as db:
        db.execute(update(Job).where(Job.id == job_id).values(execution_status=state))
        db.commit()
    source_artifact = client.get(f'/api/tools/jobs/{job_id}').json()["artifacts"][0]
    original = client.get(f'/api/tools/artifacts/{source_artifact["id"]}/download').content
    current.id = "bob"
    assert client.delete(f'/api/tools/jobs/{job_id}').status_code == 404
    current.id = "alice"
    assert client.delete(f'/api/tools/jobs/{job_id}').status_code == 204
    assert client.delete(f'/api/tools/jobs/{job_id}').status_code == 204
    assert client.get('/api/tools/jobs').json()["total"] == 1
    assert client.get('/api/tools/jobs', params={"status": state}).json()["total"] == (1 if state == "queued" else 0)
    for suffix in ["", "/result", "/issues"]:
        assert client.get(f'/api/tools/jobs/{job_id}{suffix}').status_code == 404
    for action in ["cancel", "retry"]:
        assert client.post(f'/api/tools/jobs/{job_id}/{action}').status_code == 404
    assert client.get(f'/api/tools/sources/{source["source_id"]}').status_code == 200
    assert client.get(f'/api/tools/artifacts/{source_artifact["id"]}/download').content == original
    with sessions() as db:
        deleted = db.get(Job, job_id)
        assert deleted.options_json["_deleted_at"]
        assert deleted.execution_status == ("cancelled" if state in ["queued", "running", "awaiting_input"] else state)
        assert db.get(Job, other["inspection_job_id"]).execution_status == "queued"


def test_deleted_parent_keeps_queued_revision_and_package_working(runtime):
    client, sessions, _ = runtime
    source, job = converted(runtime)
    result = client.get(f'/api/tools/jobs/{job["id"]}/result').json()
    cell = next(c for t in result["tables"] for c in t["cells"] if c["display_text"] == "000123")
    revision = client.post(f'/api/tools/jobs/{job["id"]}/revise', json={"base_revision": 1,
        "corrections": [{"target_id": cell["id"], "new_value": "000456", "reason": "verified"}]}).json()["job_id"]
    artifact = next(a for a in job["artifacts"] if a["role"] == "result")
    package = client.post('/api/tools/packages', json={"artifact_ids": [artifact["id"]], "client_request_id": jobs.uid()}).json()["job_id"]
    assert client.delete(f'/api/tools/jobs/{job["id"]}').status_code == 204
    assert pipeline.run_one(sessions)
    assert pipeline.run_one(sessions)
    assert client.get(f'/api/tools/jobs/{revision}').json()["execution_status"] == "succeeded"
    assert client.get(f'/api/tools/jobs/{package}').json()["execution_status"] == "succeeded"
    revised = client.get(f'/api/tools/jobs/{revision}/result').json()
    assert any(c["display_text"] == "000456" for t in revised["tables"] for c in t["cells"])
    with sessions() as db:
        assert db.scalar(select(Correction).where(Correction.new_job_id == revision))


def test_completed_withdraw_is_noop_and_deleted_request_cannot_resurrect(runtime):
    client, sessions, _ = runtime
    source, job = converted(runtime)
    assert client.post(f'/api/tools/jobs/{job["id"]}/cancel').json()["execution_status"] == "succeeded"
    with sessions() as db:
        request_id = db.get(Job, job["id"]).client_request_id
    assert client.delete(f'/api/tools/jobs/{job["id"]}').status_code == 204
    response = client.post('/api/tools/jobs', json={"source_id": source["source_id"],
        "operation": "word_to_excel", "options": {}, "client_request_id": request_id})
    assert response.status_code == 409 and response.json()["detail"]["code"] == "TASK_DELETED"


def test_delete_during_publication_cannot_publish_stale_output(runtime, monkeypatch):
    client, sessions, _ = runtime
    source = upload(client)
    job_id = source["inspection_job_id"]
    copytree = pipeline.shutil.copytree
    def withdraw_after_copy(*args, **kwargs):
        value = copytree(*args, **kwargs)
        assert client.delete(f'/api/tools/jobs/{job_id}').status_code == 204
        return value
    monkeypatch.setattr(pipeline.shutil, "copytree", withdraw_after_copy)
    pipeline.run_one(sessions)
    with sessions() as db:
        assert db.get(Job, job_id).execution_status == "cancelled"
        assert all(a.role == "source" for a in db.scalars(select(Artifact).where(Artifact.job_id == job_id)))
        assert db.get(Source, source["source_id"]).manifest_key == ""
