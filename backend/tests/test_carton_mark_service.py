import sys
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from pypdf import PdfWriter


BACKEND_DIR = Path(__file__).resolve().parents[1]

if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))


def _carton_mark_workbook_bytes(*values: object) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "箱唛"
    for row_number, value in enumerate(values, start=1):
        sheet.cell(row_number, 1, value)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _empty_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _use_passthrough_document_scope(monkeypatch, carton_mark) -> None:
    """Keep comparison-focused tests independent from scope-selection policy."""

    diagnostics = SimpleNamespace(describe=lambda: "测试正文选区。")
    monkeypatch.setattr(
        carton_mark,
        "prepare_carton_mark_document_scope",
        lambda excel_items, pdf_items: SimpleNamespace(
            excel_items=excel_items,
            pdf_items=pdf_items,
            requires_review=False,
            confidence=1.0,
            review_note="",
            diagnostics=diagnostics,
        ),
    )


def test_carton_mark_document_check_ignores_layout_whitespace_and_case(monkeypatch):
    from app.schemas.carton_mark import CartonMarkDocumentTextItem, CartonMarkExtractionStatus
    from app.services import carton_mark

    _use_passthrough_document_scope(monkeypatch, carton_mark)

    monkeypatch.setattr(
        carton_mark,
        "extract_pdf_document_items",
        lambda _content: carton_mark.DocumentExtraction(
            items=[CartonMarkDocumentTextItem(
                text="po no: ab-123   item: toy car",
                location="第 1 页 · 第 1 行",
            )],
            status=CartonMarkExtractionStatus(
                source="print_pdf", ok=True, engine="test", raw_text=""
            ),
            reviews=[],
        ),
    )

    result = carton_mark.build_carton_mark_document_check(
        excel_file_name="客户箱唛.xlsx",
        excel_bytes=_carton_mark_workbook_bytes("PO NO: AB-123", "ITEM: Toy Car"),
        pdf_file_name="打印箱唛.pdf",
        pdf_bytes=_empty_pdf_bytes(),
    )

    assert result.summary.overall_status == "核对通过"
    assert result.summary.pass_count == 2
    assert not result.summary.changed_count
    assert not result.summary.missing_count
    assert not result.summary.unexpected_count
    assert [item.expected_location for item in result.comparisons] == ["箱唛!A1", "箱唛!A2"]


def test_carton_mark_document_check_reports_changed_missing_and_unexpected(monkeypatch):
    from app.schemas.carton_mark import CartonMarkDocumentTextItem, CartonMarkExtractionStatus
    from app.services import carton_mark

    _use_passthrough_document_scope(monkeypatch, carton_mark)

    monkeypatch.setattr(
        carton_mark,
        "extract_pdf_document_items",
        lambda _content: carton_mark.DocumentExtraction(
            items=[
                CartonMarkDocumentTextItem(text="PO NO: AB-124", location="第 1 页 · 第 1 行"),
                CartonMarkDocumentTextItem(text="NEW WARNING", location="第 1 页 · 第 2 行"),
            ],
            status=CartonMarkExtractionStatus(
                source="print_pdf", ok=True, engine="test", raw_text=""
            ),
            reviews=[],
        ),
    )

    result = carton_mark.build_carton_mark_document_check(
        excel_file_name="客户箱唛.xlsx",
        excel_bytes=_carton_mark_workbook_bytes("PO NO: AB-123", "ITEM: Toy Car"),
        pdf_file_name="打印箱唛.pdf",
        pdf_bytes=_empty_pdf_bytes(),
    )

    assert result.summary.overall_status == "发现差异"
    assert result.summary.changed_count == 1
    assert result.summary.missing_count == 1
    assert result.summary.unexpected_count == 1
    statuses = {item.status for item in result.comparisons}
    assert statuses == {"changed", "missing", "unexpected"}
    changed = next(item for item in result.comparisons if item.status == "changed")
    assert changed.expected == "PO NO: AB-123"
    assert changed.actual == "PO NO: AB-124"
    assert changed.expected_location == "箱唛!A1"
    assert changed.actual_location == "第 1 页 · 第 1 行"


def test_carton_mark_document_check_skips_hidden_cells_and_reviews_uncached_formula(monkeypatch):
    from app.schemas.carton_mark import CartonMarkDocumentTextItem, CartonMarkExtractionStatus
    from app.services import carton_mark

    _use_passthrough_document_scope(monkeypatch, carton_mark)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "箱唛"
    sheet["A1"] = "VISIBLE"
    sheet["A2"] = "HIDDEN ROW"
    sheet.row_dimensions[2].hidden = True
    sheet["B1"] = "HIDDEN COL"
    sheet.column_dimensions["B"].hidden = True
    sheet["A3"] = '=CONCAT("PO", "123")'
    output = BytesIO()
    workbook.save(output)

    monkeypatch.setattr(
        carton_mark,
        "extract_pdf_document_items",
        lambda _content: carton_mark.DocumentExtraction(
            items=[CartonMarkDocumentTextItem(text="VISIBLE", location="第 1 页 · 第 1 行")],
            status=CartonMarkExtractionStatus(
                source="print_pdf", ok=True, engine="test", raw_text=""
            ),
            reviews=[],
        ),
    )

    result = carton_mark.build_carton_mark_document_check(
        excel_file_name="客户箱唛.xlsm",
        excel_bytes=output.getvalue(),
        pdf_file_name="打印箱唛.pdf",
        pdf_bytes=_empty_pdf_bytes(),
    )

    assert [item.text for item in result.excel_items] == ["VISIBLE"]
    assert result.summary.overall_status == "需复核"
    assert result.summary.review_count == 1
    review = next(item for item in result.comparisons if item.status == "review")
    assert review.expected_location == "箱唛!A3"
    assert "公式" in review.note


def test_excel_document_extraction_resolves_direct_reference_without_calculating_formula():
    from app.services.carton_mark import extract_excel_document_items

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "箱唛"
    sheet["A1"] = "ITEM 100372"
    sheet["A2"] = "=A1"
    sheet["A3"] = '=CONCAT("ITEM", " 100372")'
    output = BytesIO()
    workbook.save(output)

    extraction = extract_excel_document_items("客户箱唛.xlsx", output.getvalue())

    assert [item.text for item in extraction.items] == ["ITEM 100372", "ITEM 100372"]
    assert [location for location, _ in extraction.reviews] == ["箱唛!A3"]


def test_excel_document_extraction_skips_intermediate_direct_reference_control():
    from app.services.carton_mark import extract_excel_document_items

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "箱唛"
    sheet["A1"] = "ITEM 100372"
    sheet["A2"] = "=A1"
    sheet["A3"] = "=A2"
    output = BytesIO()
    workbook.save(output)

    extraction = extract_excel_document_items("客户箱唛.xlsx", output.getvalue())

    assert [(item.location, item.text) for item in extraction.items] == [
        ("箱唛!A1", "ITEM 100372"),
        ("箱唛!A3", "ITEM 100372"),
    ]
    assert not extraction.reviews


def test_excel_document_extraction_reviews_overlong_direct_reference_chain():
    from app.services.carton_mark import extract_excel_document_items

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "箱唛"
    sheet["A1"] = "ITEM 100372"
    for row in range(2, 132):
        sheet.cell(row, 1, f"=A{row - 1}")
    output = BytesIO()
    workbook.save(output)

    extraction = extract_excel_document_items("客户箱唛.xlsx", output.getvalue())

    assert [(item.location, item.text) for item in extraction.items] == [
        ("箱唛!A1", "ITEM 100372"),
    ]
    assert [location for location, _note in extraction.reviews] == [
        "箱唛!A130",
        "箱唛!A131",
    ]


@pytest.mark.parametrize(
    ("file_name", "content", "kind", "message"),
    [
        ("contract.pdf", b"not excel", "excel", "文件类型无效"),
        ("contract.xlsx", b"not a zip", "excel", "内容与 OOXML 扩展名不符"),
        ("contract.xls", b"not old excel", "excel", "不是有效的旧版 .xls"),
        ("print.pdf", b"not pdf", "pdf", "内容与 .pdf 扩展名不符"),
    ],
)
def test_carton_mark_document_file_validation_rejects_spoofed_content(
    file_name: str,
    content: bytes,
    kind: str,
    message: str,
):
    from app.services.carton_mark import CartonMarkDocumentError, validate_carton_mark_document_file

    with pytest.raises(CartonMarkDocumentError, match=message):
        validate_carton_mark_document_file(file_name, content, kind=kind)


def test_pdf_document_reader_uses_lenient_mode(monkeypatch):
    import pypdf

    from app.services.carton_mark import extract_pdf_document_items

    strict_values: list[bool] = []

    class FakePage:
        @staticmethod
        def extract_text() -> str:
            return "PO 123"

    class FakeReader:
        is_encrypted = False
        pages = [FakePage()]

        def __init__(self, _stream, *, strict: bool):
            strict_values.append(strict)

    monkeypatch.setattr(pypdf, "PdfReader", FakeReader)

    extraction = extract_pdf_document_items(b"%PDF-1.7\n")

    assert strict_values == [False]
    assert [item.text for item in extraction.items] == ["PO 123"]


def test_document_comparison_keeps_identical_repetition_on_another_page():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import (
        compare_document_text_items,
        summarize_document_comparisons,
    )

    comparisons = compare_document_text_items(
        [CartonMarkDocumentTextItem(text="PO 123", location="箱唛!A1")],
        [
            CartonMarkDocumentTextItem(text="PO 123", location="第 1 页 · 第 1 行"),
            CartonMarkDocumentTextItem(text="PO 123", location="第 2 页 · 第 1 行"),
        ],
    )

    assert [item.status for item in comparisons] == ["pass", "unexpected"]
    assert summarize_document_comparisons(comparisons).overall_status == "发现差异"


def test_document_comparison_ignores_identical_front_and_side_panels_on_same_page():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [CartonMarkDocumentTextItem(text="PO 123", location="箱唛!A1")],
        [
            CartonMarkDocumentTextItem(text="PO 123", location="第 1 页 · 正唛 · 第 1 行"),
            CartonMarkDocumentTextItem(text="PO 123", location="第 1 页 · 侧唛 · 第 1 行"),
        ],
    )

    assert [item.status for item in comparisons] == ["pass"]


def test_document_comparison_ignores_residual_layout_separators():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import (
        compare_document_text_items,
        summarize_document_comparisons,
    )

    comparisons = compare_document_text_items(
        [
            CartonMarkDocumentTextItem(text="PO", location="箱唛!A1"),
            CartonMarkDocumentTextItem(text="123", location="箱唛!B1", field_key="po"),
        ],
        [CartonMarkDocumentTextItem(
            text="PO: | • 123",
            location="第 1 页 · 第 1 行",
        )],
    )

    assert [item.status for item in comparisons] == ["pass", "pass"]
    assert summarize_document_comparisons(comparisons).overall_status == "核对通过"


def test_document_comparison_ignores_split_labels_and_known_repeated_values():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [
            CartonMarkDocumentTextItem(text="415647", location="箱唛!A1", field_key="item"),
            CartonMarkDocumentTextItem(text="ZURU.INC", location="箱唛!A2"),
        ],
        [
            CartonMarkDocumentTextItem(text="ITEM NO.: 415647", location="第 1 页 · 正唛 · 第 1 行"),
            CartonMarkDocumentTextItem(text="NAME OF SUPPLIER: ZURU.INC", location="第 1 页 · 正唛 · 第 2 行"),
            CartonMarkDocumentTextItem(text="ITEM NO.: 415647", location="第 1 页 · 侧唛 · 第 1 行"),
        ],
    )

    assert [item.status for item in comparisons] == ["pass", "pass"]


def test_document_comparison_ignores_empty_physical_placeholders_but_keeps_values():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [CartonMarkDocumentTextItem(text="ITEM 92121", location="箱唛!A1")],
        [
            CartonMarkDocumentTextItem(text="ITEM 92121", location="第 1 页 · 第 1 行"),
            CartonMarkDocumentTextItem(text="W.G.R.: KG", location="第 1 页 · 第 2 行"),
            CartonMarkDocumentTextItem(text="W.NET: KG", location="第 1 页 · 第 3 行"),
            CartonMarkDocumentTextItem(text="MEASUREMENT: x x CM", location="第 1 页 · 第 4 行"),
            CartonMarkDocumentTextItem(text="QUANTITY: 12 PCS", location="第 1 页 · 第 5 行"),
            CartonMarkDocumentTextItem(text="CTN QTY: 42 CTNS", location="第 1 页 · 第 6 行"),
        ],
    )

    unexpected = [item.actual for item in comparisons if item.status == "unexpected"]
    assert unexpected == ["QUANTITY: 12 PCS", "CTN QTY: 42 CTNS"]


def test_document_comparison_ignores_pdf_only_text_inside_graphics():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [CartonMarkDocumentTextItem(text="ITEM NO: 7044498", location="箱唛!A1")],
        [
            CartonMarkDocumentTextItem(
                text="ITEM NO: 7044498",
                location="第 1 页 · 正唛 · 第 1 行",
            ),
            CartonMarkDocumentTextItem(
                text="FTC",
                location="第 1 页 · 正唛 · 第 2 行",
                is_graphic_text=True,
            ),
            CartonMarkDocumentTextItem(
                text="New Zealand",
                location="第 1 页 · 正唛 · 第 3 行",
                is_graphic_text=True,
            ),
            CartonMarkDocumentTextItem(
                text="NEW WARNING",
                location="第 1 页 · 正唛 · 第 4 行",
            ),
        ],
    )

    assert [item.status for item in comparisons] == ["pass", "unexpected"]
    assert comparisons[-1].actual == "NEW WARNING"


def test_document_comparison_does_not_fuzzy_match_unlabelled_numbers():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [CartonMarkDocumentTextItem(text="3.12", location="箱唛!A1")],
        [CartonMarkDocumentTextItem(text="12", location="第 1 页 · 第 1 行")],
    )

    assert [item.status for item in comparisons] == ["missing", "unexpected"]


def test_document_comparison_treats_underscore_placeholder_as_layout():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [CartonMarkDocumentTextItem(text="_OF", location="箱唛!A1")],
        [CartonMarkDocumentTextItem(text="OF", location="第 1 页 · 第 1 行")],
    )

    assert [item.status for item in comparisons] == ["pass"]


def test_document_comparison_does_not_consume_short_quantity_inside_item_number():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [
            CartonMarkDocumentTextItem(text="9560", location="箱唛!A1", field_key="item"),
            CartonMarkDocumentTextItem(text="6", location="箱唛!B1", field_key="quantity"),
        ],
        [
            CartonMarkDocumentTextItem(text="ARTICLE NO.: 9560", location="第 1 页 · 正唛 · 第 1 行"),
            CartonMarkDocumentTextItem(text="ARTICLE NO.: 9560", location="第 1 页 · 侧唛 · 第 1 行"),
            CartonMarkDocumentTextItem(text="NUMBER OF PIECES: 6 PCS", location="第 1 页 · 正唛 · 第 2 行"),
        ],
    )

    assert [item.status for item in comparisons] == ["pass", "pass"]


def test_document_comparison_keeps_same_number_added_under_different_field():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [CartonMarkDocumentTextItem(
            text="12 PCS",
            location="箱唛!B2",
            field_key="quantity",
        )],
        [
            CartonMarkDocumentTextItem(
                text="QTY: 12 PCS",
                location="第 1 页 · 第 1 行",
            ),
            CartonMarkDocumentTextItem(
                text="GROSS WEIGHT: 12 KG",
                location="第 1 页 · 第 2 行",
            ),
        ],
    )

    assert [item.status for item in comparisons] == ["pass", "unexpected"]
    assert comparisons[-1].actual == "GROSS WEIGHT: 12 KG"


def test_document_comparison_does_not_match_known_quantity_to_bare_unknown_number():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [CartonMarkDocumentTextItem(
            text="12",
            location="箱唛!B2",
            field_key="quantity",
        )],
        [CartonMarkDocumentTextItem(text="12", location="第 1 页 · 第 1 行")],
    )

    assert [item.status for item in comparisons] == ["missing", "unexpected"]


def test_document_comparison_resolves_multiple_fields_inside_one_pdf_line():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [
            CartonMarkDocumentTextItem(text="AB-123", location="箱唛!A1", field_key="po"),
            CartonMarkDocumentTextItem(text="TOY CAR", location="箱唛!B1", field_key="item"),
        ],
        [CartonMarkDocumentTextItem(
            text="PO NO: AB-123   ITEM: TOY CAR",
            location="第 1 页 · 第 1 行",
        )],
    )

    assert [item.status for item in comparisons] == ["pass", "pass"]


def test_document_comparison_resolves_label_inclusive_fields_inside_one_pdf_line():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [
            CartonMarkDocumentTextItem(text="PO NO: AB-123", location="箱唛!A1", field_key="po"),
            CartonMarkDocumentTextItem(text="ITEM: TOY CAR", location="箱唛!B1", field_key="item"),
        ],
        [CartonMarkDocumentTextItem(
            text="PO NO: AB-123   ITEM: TOY CAR",
            location="第 1 页 · 第 1 行",
        )],
    )

    assert [item.status for item in comparisons] == ["pass", "pass"]


def test_document_comparison_keeps_extra_short_quantity_as_unexpected():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [CartonMarkDocumentTextItem(
            text="12 PCS",
            location="箱唛!A1",
            field_key="quantity",
        )],
        [
            CartonMarkDocumentTextItem(text="QTY: 12 PCS", location="第 1 页 · 第 1 行"),
            CartonMarkDocumentTextItem(text="QTY: 2 PCS", location="第 1 页 · 第 2 行"),
        ],
    )

    assert [item.status for item in comparisons] == ["pass", "unexpected"]
    assert comparisons[-1].actual == "QTY: 2 PCS"


def test_document_comparison_matches_label_only_source_to_combined_pdf_line():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [CartonMarkDocumentTextItem(
            text="CTN QTY:",
            location="箱唛!A1",
            field_key="carton_no",
        )],
        [CartonMarkDocumentTextItem(
            text="CTN QTY: / 42 CTNS",
            location="第 1 页 · 第 1 行",
            field_key="quantity",
        )],
    )

    assert [item.status for item in comparisons] == ["pass", "unexpected"]
    assert "42" in comparisons[-1].actual


def test_document_comparison_matches_long_identifier_in_slash_layout_context():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [
            CartonMarkDocumentTextItem(text="FTC", location="箱唛!A1"),
            CartonMarkDocumentTextItem(text="D685", location="箱唛!B1"),
            CartonMarkDocumentTextItem(text="SSF248", location="箱唛!C1"),
            CartonMarkDocumentTextItem(
                text="4504244551",
                location="箱唛!D1",
                field_key="po",
            ),
        ],
        [CartonMarkDocumentTextItem(
            text="FTC / D685 / SSF248 / 4504244551",
            location="第 1 页 · 第 1 行",
        )],
    )

    assert [item.status for item in comparisons] == ["pass"] * 4


def test_document_comparison_does_not_match_identifier_inside_warning_text():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [CartonMarkDocumentTextItem(
            text="123456",
            location="箱唛!A1",
            field_key="item",
        )],
        [CartonMarkDocumentTextItem(
            text="NEW WARNING 123456",
            location="第 1 页 · 第 1 行",
        )],
    )

    assert [item.status for item in comparisons] == ["missing", "unexpected"]


def test_document_comparison_does_not_match_unknown_bare_number_to_weight():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [CartonMarkDocumentTextItem(text="12", location="箱唛!B2")],
        [CartonMarkDocumentTextItem(
            text="GROSS WEIGHT: 12 KG",
            location="第 1 页 · 第 1 行",
        )],
    )

    assert [item.status for item in comparisons] == ["missing", "unexpected"]


def test_document_comparison_does_not_auto_pass_two_unknown_short_numbers():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [CartonMarkDocumentTextItem(text="12", location="箱唛!B2")],
        [CartonMarkDocumentTextItem(text="12", location="第 1 页 · 第 1 行")],
    )

    assert [item.status for item in comparisons] == ["missing", "unexpected"]


def test_document_comparison_keeps_short_quantity_repeated_across_pages():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import _append_pdf_lines, compare_document_text_items

    actual_items: list[CartonMarkDocumentTextItem] = []
    _append_pdf_lines(actual_items, "QUANTITY:\n12\nPCS", 1)
    _append_pdf_lines(actual_items, "QUANTITY:\n12\nPCS", 2)

    assert [item.field_key for item in actual_items] == [
        "quantity", "quantity", "quantity",
        "quantity", "quantity", "quantity",
    ]

    comparisons = compare_document_text_items(
        [],
        actual_items,
    )

    assert any(item.status == "unexpected" and item.actual == "12" for item in comparisons)


def test_pdf_pending_field_is_consumed_by_only_the_next_value():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import _append_pdf_lines

    actual_items: list[CartonMarkDocumentTextItem] = []
    _append_pdf_lines(actual_items, "GROSS WEIGHT:\n5.04\n99", 1)

    assert [item.field_key for item in actual_items] == ["gw", "gw", ""]


def test_pdf_inline_field_does_not_start_pending_context():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import _append_pdf_lines

    actual_items: list[CartonMarkDocumentTextItem] = []
    _append_pdf_lines(actual_items, "QTY: 6 PCS\n5.04 KGS", 1)

    assert [item.field_key for item in actual_items] == ["quantity", ""]


def test_pdf_reverse_order_value_column_pairs_with_following_labels():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import _append_pdf_lines

    actual_items: list[CartonMarkDocumentTextItem] = []
    _append_pdf_lines(
        actual_items,
        "ITEM DESCRIPTION:\nTOY SERIES 1\n4PCS\n82\nMADE IN CHINA\nCTN QTY:\nCTN NUMBER: /",
        1,
    )

    by_text = {item.text: item.field_key for item in actual_items}
    assert by_text["4PCS"] == "quantity"
    assert by_text["82"] == "carton_no"


def test_pdf_description_continuation_does_not_infer_field_from_internal_color_word():
    from app.services.carton_mark import _document_field_key

    assert _document_field_key("-MOMMY TURTLE SURPRISE PLAYSET-COLOR BOX-4PCS/CTN") == ""
    assert _document_field_key("& PONIES-Sparkle Unicorn Styling Open Box,6pcs/CTN") == ""


def test_pdf_spaced_barcode_digits_have_barcode_context():
    from app.services.carton_mark import _document_field_key

    assert _document_field_key("1 0 1 9 3 0 5 2 0 7 8 4 0 4") == "barcode"


def test_document_comparison_keeps_independent_substring_addition():
    from app.schemas.carton_mark import CartonMarkDocumentTextItem
    from app.services.carton_mark import compare_document_text_items

    comparisons = compare_document_text_items(
        [CartonMarkDocumentTextItem(
            text="WARNING BATTERY INCLUDED",
            location="箱唛!A1",
        )],
        [
            CartonMarkDocumentTextItem(
                text="WARNING BATTERY INCLUDED",
                location="第 1 页 · 第 1 行",
            ),
            CartonMarkDocumentTextItem(
                text="BATTERY INCLUDED",
                location="第 1 页 · 第 2 行",
            ),
        ],
    )

    assert [item.status for item in comparisons] == ["pass", "unexpected"]


@pytest.mark.parametrize(
    ("engine", "ok", "reviews"),
    [
        ("pytesseract", True, []),
        ("pypdf", False, [("第 2 页", "该页没有可读取文字层，请人工确认。")]),
    ],
    ids=["ocr", "incomplete-text-layer"],
)
def test_document_check_caps_unreliable_pdf_differences_at_review(
    monkeypatch,
    engine: str,
    ok: bool,
    reviews: list[tuple[str, str]],
):
    from app.schemas.carton_mark import CartonMarkDocumentTextItem, CartonMarkExtractionStatus
    from app.services import carton_mark

    _use_passthrough_document_scope(monkeypatch, carton_mark)

    monkeypatch.setattr(
        carton_mark,
        "extract_pdf_document_items",
        lambda _content: carton_mark.DocumentExtraction(
            items=[CartonMarkDocumentTextItem(
                text="PO NO: AB-124",
                location="第 1 页 · 第 1 行",
            )],
            status=CartonMarkExtractionStatus(
                source="print_pdf",
                ok=ok,
                engine=engine,
                message="PDF 提取结果需要人工复核。",
                raw_text="",
            ),
            reviews=reviews,
        ),
    )

    result = carton_mark.build_carton_mark_document_check(
        excel_file_name="客户箱唛.xlsx",
        excel_bytes=_carton_mark_workbook_bytes("PO NO: AB-123"),
        pdf_file_name="打印箱唛.pdf",
        pdf_bytes=_empty_pdf_bytes(),
    )

    assert result.summary.changed_count == 1
    assert result.summary.overall_status == "需复核"


def test_document_check_scope_review_forces_manual_review(monkeypatch):
    from app.schemas.carton_mark import CartonMarkDocumentTextItem, CartonMarkExtractionStatus
    from app.services import carton_mark

    monkeypatch.setattr(
        carton_mark,
        "extract_pdf_document_items",
        lambda _content: carton_mark.DocumentExtraction(
            items=[CartonMarkDocumentTextItem(
                text="PO NO: AB-124",
                location="第 1 页 · 正唛 · 第 1 行",
            )],
            status=CartonMarkExtractionStatus(
                source="print_pdf", ok=True, engine="pypdf-vector-coordinates", raw_text=""
            ),
            reviews=[],
        ),
    )
    diagnostics = SimpleNamespace(describe=lambda: "候选正文有两个相近选区。")
    monkeypatch.setattr(
        carton_mark,
        "prepare_carton_mark_document_scope",
        lambda excel_items, pdf_items: SimpleNamespace(
            excel_items=excel_items,
            pdf_items=pdf_items,
            requires_review=True,
            confidence=0.51,
            review_note="正文选区证据不足，本次结果必须人工复核。",
            diagnostics=diagnostics,
        ),
    )

    result = carton_mark.build_carton_mark_document_check(
        excel_file_name="客户箱唛.xlsx",
        excel_bytes=_carton_mark_workbook_bytes("PO NO: AB-123"),
        pdf_file_name="打印箱唛.pdf",
        pdf_bytes=_empty_pdf_bytes(),
    )

    assert result.summary.changed_count == 1
    assert result.summary.overall_status == "需复核"
    assert any(item.status == "review" and "正文选区" in item.note for item in result.comparisons)
    scope_status = next(status for status in result.extraction if status.source == "document_scope")
    assert not scope_status.ok
    assert scope_status.requires_review
    assert scope_status.match_confidence == pytest.approx(0.51)
    assert "人工复核" in scope_status.message


def test_document_check_high_confidence_scope_keeps_real_missing_values(monkeypatch):
    from app.schemas.carton_mark import CartonMarkDocumentTextItem, CartonMarkExtractionStatus
    from app.services import carton_mark

    _use_passthrough_document_scope(monkeypatch, carton_mark)
    monkeypatch.setattr(
        carton_mark,
        "extract_pdf_document_items",
        lambda _content: carton_mark.DocumentExtraction(
            items=[CartonMarkDocumentTextItem(
                text="ITEM 92121TQ1",
                location="第 1 页 · 侧唛 · 第 1 行",
            )],
            status=CartonMarkExtractionStatus(
                source="print_pdf", ok=True, engine="pypdf-vector-coordinates", raw_text=""
            ),
            reviews=[],
        ),
    )

    result = carton_mark.build_carton_mark_document_check(
        excel_file_name="客户箱唛.xlsx",
        excel_bytes=_carton_mark_workbook_bytes("ITEM 92121TQ1", "3.74", "3.14", "55.6 x 19.7 x 57.5"),
        pdf_file_name="打印箱唛.pdf",
        pdf_bytes=_empty_pdf_bytes(),
    )

    assert result.summary.overall_status == "发现差异"
    assert result.summary.missing_count == 3
    assert {
        item.expected
        for item in result.comparisons
        if item.status == "missing"
    } == {"3.74", "3.14", "55.6 x 19.7 x 57.5"}


def test_pdf_document_line_filter_drops_broken_font_noise():
    from app.services.carton_mark import _append_pdf_lines

    items = []
    _append_pdf_lines(
        items,
        "PO NUMBER: 484860\nѓᆽԄժđοᆼདൌ࠽Ԅժοေ౰טᆜႆ඗đ\nQUANTITY: 12 PCS",
        1,
    )

    assert [item.text for item in items] == ["PO NUMBER: 484860", "QUANTITY: 12 PCS"]


def test_pdf_document_line_filter_drops_leading_repeated_artwork_glyphs_only():
    from app.services.carton_mark import _append_pdf_lines

    items = []
    _append_pdf_lines(items, "9 9\nITEM NO: 9574\n9 9", 1)

    assert [item.text for item in items] == ["ITEM NO: 9574", "9 9"]


@pytest.mark.parametrize("noise", ["ο DNಀႆ Ġ", "ZURUᆜႆ đ", "ᆜႆ"])
def test_pdf_document_line_filter_drops_short_mixed_script_font_noise(noise: str):
    from app.services.carton_mark import _clean_pdf_document_line

    assert _clean_pdf_document_line(noise) == ""


@pytest.mark.parametrize(
    ("label", "field_key"),
    [
        ("PI NO.", "pi_no"),
        ("G.W.: KGS", "gw"),
        ("W.NET: KG", "nw"),
        ("Qty/Ctn:", "quantity"),
        ("Carton Size(cm):", "measurement"),
        ("CTN QTY:", "quantity"),
        ("BARCODE NUMBER:", "barcode"),
        ("GROSS WT: KG", "gw"),
        ("NET WT: KG", "nw"),
    ],
)
def test_pdf_vector_label_info_normalizes_label_punctuation(label: str, field_key: str):
    from app.services.carton_mark import _pdf_vector_label_info

    label_only, actual_key, starts_label = _pdf_vector_label_info(label)

    assert starts_label
    assert label_only
    assert actual_key == field_key


def test_pdf_vector_document_scope_preserves_repeated_multi_page_bodies_and_additions():
    from app.services.carton_mark import PdfTextSpan, build_pdf_vector_document_items

    def page(page_index: int, item: str):
        return build_pdf_vector_document_items(
            [
                PdfTextSpan(text=f"ITEM NO: {item}", x=120, y=280),
                PdfTextSpan(text="DESCRIPTION: FUGGLER BADDIE", x=120, y=240),
                PdfTextSpan(text="MADE IN CHINA", x=120, y=200),
                PdfTextSpan(text=f"ITEM NO: {item}", x=540, y=280),
                PdfTextSpan(text="QUANTITY: 12 PCS", x=540, y=220),
                PdfTextSpan(text="CTN QTY: 42 CTNS", x=540, y=180),
                PdfTextSpan(text="PRODUCTION HEADER MUST NOT COMPARE", x=20, y=560),
                PdfTextSpan(text="ѓᆽԄժđοᆼདൌ", x=420, y=300),
            ],
            page_width=842,
            page_height=595,
            page_index=page_index,
        )

    items = [*page(1, "415647"), *page(2, "415648")]

    assert any(item.text == "ITEM NO: 415647" and item.location.startswith("第 1 页") for item in items)
    assert any(item.text == "ITEM NO: 415648" and item.location.startswith("第 2 页") for item in items)
    assert sum(item.text == "QUANTITY: 12 PCS" for item in items) == 2
    assert sum(item.text == "CTN QTY: 42 CTNS" for item in items) == 2
    assert all("PRODUCTION HEADER" not in item.text for item in items)
    assert all("ѓ" not in item.text for item in items)


def test_pdf_vector_extraction_marks_non_rectangular_artwork_text_only():
    from app.services.carton_mark import _extract_pdf_page_vector_document_items

    class MediaBox:
        width = 842
        height = 595

    class Page:
        mediabox = MediaBox()

        def extract_text(self, *, visitor_text, visitor_operand_before):
            identity = [1, 0, 0, 1, 0, 0]
            visitor_operand_before(b"m", [100, 250], identity, identity)
            visitor_operand_before(b"l", [160, 190], identity, identity)
            visitor_operand_before(b"S", [], identity, identity)
            visitor_operand_before(b"re", [100, 120, 180, 45], identity, identity)
            visitor_operand_before(b"S", [], identity, identity)
            for text, x, y in (
                ("FTC", 130, 220),
                ("New Zealand", 120, 200),
                ("ITEM NO: 7044498", 120, 155),
                ("QTY: 8 PCS", 120, 135),
                ("NEW WARNING", 120, 105),
            ):
                visitor_text(text, identity, [1, 0, 0, 1, x, y], None, 12)
            return ""

    items = _extract_pdf_page_vector_document_items(Page(), 1)
    by_text = {item.text: item for item in items}

    assert by_text["FTC"].is_graphic_text
    assert by_text["New Zealand"].is_graphic_text
    assert not by_text["ITEM NO: 7044498"].is_graphic_text
    assert not by_text["QTY: 8 PCS"].is_graphic_text
    assert not by_text["NEW WARNING"].is_graphic_text


def test_pdf_vector_document_fields_pair_same_visual_rows():
    from app.services.carton_mark import PdfTextSpan, build_pdf_vector_document_items

    items = build_pdf_vector_document_items(
        [
            PdfTextSpan(text="ITEM: 77770GQ2", x=120, y=250),
            PdfTextSpan(text="CTN QTY:", x=120, y=210),
            PdfTextSpan(text="CTN NUMBER:", x=120, y=195),
            PdfTextSpan(text="50 PCS", x=250, y=210),
            PdfTextSpan(text="49", x=250, y=195),
            PdfTextSpan(text="MADE IN CHINA", x=120, y=180),
        ],
        page_width=842,
        page_height=595,
        page_index=1,
    )

    by_text = {item.text: item.field_key for item in items}
    assert by_text["50 PCS"] == "quantity"
    assert by_text["49"] == "carton_no"


def test_pdf_vector_document_fields_pair_label_then_value_columns():
    from app.services.carton_mark import PdfTextSpan, build_pdf_vector_document_items

    items = build_pdf_vector_document_items(
        [
            PdfTextSpan(text="ITEM NO:", x=550, y=326),
            PdfTextSpan(text="BARCODE:", x=550, y=268),
            PdfTextSpan(text="BRAND:", x=550, y=255),
            PdfTextSpan(text="QTY/CTN:", x=550, y=242),
            PdfTextSpan(text="CARTON NET WEIGHT:", x=550, y=229),
            PdfTextSpan(text="CARTON GROSS WEIGHT:", x=550, y=216),
            PdfTextSpan(text="CARTON SIZE(CM):", x=550, y=203),
            PdfTextSpan(text="CARTON CBM:", x=550, y=190),
            PdfTextSpan(text="PI No.:", x=550, y=177),
            PdfTextSpan(text="9531", x=632, y=326),
            PdfTextSpan(text="193052057565", x=632, y=267),
            PdfTextSpan(text="ZURU PETS ALIVE", x=632, y=132),
            PdfTextSpan(text="6 PCS", x=632, y=119),
            PdfTextSpan(text="5.04 KGS", x=632, y=106),
            PdfTextSpan(text="5.93 KGS", x=632, y=-30),
            PdfTextSpan(text="56.5 X 31.4 X 63.5 CM", x=632, y=-43),
            PdfTextSpan(text="0.113 CBM", x=632, y=-180),
            PdfTextSpan(text="PI-2025-001", x=632, y=-317),
        ],
        page_width=842,
        page_height=595,
        page_index=1,
    )

    by_text = {item.text: item.field_key for item in items}
    assert by_text["9531"] == "item"
    assert by_text["193052057565"] == "barcode"
    assert by_text["ZURU PETS ALIVE"] == ""
    assert by_text["6 PCS"] == "quantity"
    assert by_text["5.04 KGS"] == "nw"
    assert by_text["5.93 KGS"] == "gw"
    assert by_text["56.5 X 31.4 X 63.5 CM"] == "measurement"
    assert by_text["0.113 CBM"] == ""
    assert by_text["PI-2025-001"] == "pi_no"


def test_pdf_ocr_configuration_error_is_not_downgraded_to_document_review(monkeypatch):
    from app.services import carton_mark

    class FakeRenderedPage:
        @staticmethod
        def to_pil():
            return FakeImage()

    class FakeImage:
        def convert(self, mode: str):
            assert mode == "RGB"
            return self

    class FakePage:
        @staticmethod
        def render(*, scale: int):
            assert scale == 2
            return FakeRenderedPage()

    class FakeDocument:
        @staticmethod
        def __getitem__(_page_index: int):
            return FakePage()

    fake_pdfium = SimpleNamespace(PdfDocument=lambda _content: FakeDocument())
    monkeypatch.setitem(sys.modules, "pypdfium2", fake_pdfium)

    def raise_configuration_error(_image):
        raise carton_mark.CartonMarkDocumentConfigurationError(
            "服务器未配置 OCR 组件"
        )

    monkeypatch.setattr(carton_mark, "_ocr_document_page", raise_configuration_error)

    with pytest.raises(
        carton_mark.CartonMarkDocumentConfigurationError,
        match="服务器未配置 OCR 组件",
    ):
        carton_mark._extract_pdf_document_items_with_ocr(b"%PDF-1.7\n", 1)


def test_carton_mark_auto_check_extracts_template_fields_and_requests_photo_review():
    from app.services.carton_mark import build_carton_mark_auto_check

    response = build_carton_mark_auto_check(
        pdf_bytes=(
            b"SHIPPING IDENTIFICATION MARK\n"
            b"NUMERO DE PEDIDO 203302017\n"
            b"ITEM 700142617\n"
            b"G.W. 4.3 KGS\n"
            b"N.W. 3.1 KGS\n"
            b"BARCODE 197919619\n"
        ),
        front_image_bytes=b"not-an-image",
        side_image_bytes=b"not-an-image",
        customer_name="Dickie",
        po="203302017",
        item="700142617",
    )

    template_fields = {field.key: field.value for field in response.template_fields}
    front_template_fields = {field.key: field.value for field in response.front_template_fields}
    side_template_fields = {field.key: field.value for field in response.side_template_fields}
    assert template_fields["customer_name"] == "Dickie"
    assert template_fields["po"] == "203302017"
    assert template_fields["item"] == "700142617"
    assert template_fields["barcode"] == "197919619"
    assert front_template_fields["po"] == "203302017"
    assert side_template_fields["item"] == "700142617"

    assert response.summary.overall_status == "需复核"
    assert response.summary.review_count > 0
    assert any(item.side == "front" and item.field_key == "po" and item.status == "review" for item in response.comparisons)
    assert any(status.source == "front_photo" and not status.ok for status in response.extraction)


def test_carton_mark_batch_auto_check_keeps_front_and_side_images_independent(monkeypatch):
    from app.schemas.carton_mark import CartonMarkExtractionStatus
    from app.services import carton_mark

    def status(source: str) -> CartonMarkExtractionStatus:
        return CartonMarkExtractionStatus(
            source=source,
            ok=True,
            engine="test",
            raw_text="",
        )

    monkeypatch.setattr(carton_mark, "extract_pdf_text", lambda _: (
        "PO 62098330\nITEM 203302017\nCOLOR MULTICOLOR",
        status("pdf_template"),
    ))
    monkeypatch.setattr(carton_mark, "extract_pdf_template_side_texts", lambda *_args, **_kwargs: (
        "PO 62098330\nITEM 203302017\nCOLOR MULTICOLOR",
        "PO 62098330\nITEM 203302017\nQTY 24 PCS",
        status("pdf_layout"),
    ))
    monkeypatch.setattr(carton_mark, "extract_image_text", lambda image, *, source: (
        image.decode("utf-8"),
        status(source),
    ))

    response = carton_mark.build_carton_mark_batch_auto_check(
        pdf_bytes=b"pdf",
        front_images=[("front-01.jpg", b"PO 62098330\nITEM 203302017\nCOLOR MULTICOLOR")],
        side_images=[("side-01.jpg", b"PO 62098330\nITEM 203302017\nQTY 24 PCS")],
        customer_name="Dickie",
        item="203302017",
    )

    assert [(item.side, item.file_name, item.file_index) for item in response.items] == [
        ("front", "front-01.jpg", 0),
        ("side", "side-01.jpg", 0),
    ]
    assert {item.side for item in response.items[0].result.comparisons} == {"front"}
    assert {item.side for item in response.items[1].result.comparisons} == {"side"}


def test_carton_mark_field_extraction_handles_noisy_photo_ocr_rows():
    from app.services.carton_mark import extract_fields

    fields = {
        field.key: field.value
        for field in extract_fields(
            (
                "P.0 N0 | 2033 | ITEM N0 | 2017\n"
                "QTY/CTN 24 PCS\n"
                "G W 4.3 KGS\n"
                "N W 3.1 KGS\n"
                "BULTO 1-10\n"
            ),
            source="front_photo",
        )
    }

    assert fields["po"] == "2033"
    assert fields["item"] == "2017"
    assert fields["quantity"] == "24"
    assert fields["gw"] == "4.3 KGS"
    assert fields["nw"] == "3.1 KGS"
    assert fields["carton_no"] == "1-10"


def test_carton_mark_photo_ocr_variants_include_detected_label_crop():
    from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageOps

    from app.services.carton_mark import build_photo_ocr_variants, locate_photo_mark_boxes

    image = Image.new("RGB", (1200, 900), (174, 142, 96))
    draw = ImageDraw.Draw(image)
    draw.rectangle((270, 170, 910, 730), fill=(226, 214, 176), outline=(45, 42, 34), width=4)
    for y in range(245, 705, 70):
        draw.line((285, y, 895, y), fill=(45, 42, 34), width=3)
    for x in (470, 690):
        draw.line((x, 185, x, 715), fill=(45, 42, 34), width=3)
    draw.text((305, 205), "PO NO 2033", fill=(20, 20, 20))
    draw.text((305, 275), "ITEM NO 2017", fill=(20, 20, 20))
    draw.text((305, 345), "QTY/CTN 24 PCS", fill=(20, 20, 20))
    draw.text((305, 415), "G.W 4.3 KGS", fill=(20, 20, 20))
    draw.rectangle((285, 660, 895, 705), fill=(96, 24, 36))

    boxes = locate_photo_mark_boxes(image, ImageFilter, ImageOps)
    variants = build_photo_ocr_variants(image, ImageEnhance, ImageFilter, ImageOps)
    largest_box = boxes[0]

    assert boxes
    assert largest_box[0] > 0
    assert largest_box[1] > 0
    assert largest_box[2] < image.width
    assert largest_box[3] < image.height
    assert len(variants) > 3


def test_carton_mark_ocr_word_rows_recover_table_label_values():
    from app.services.carton_mark import OcrWord, build_ocr_word_table_text, extract_fields

    words = [
        OcrWord(text="NUMERO", left=10, top=10, right=70, bottom=28, confidence=82),
        OcrWord(text="DE", left=76, top=10, right=98, bottom=28, confidence=82),
        OcrWord(text="PEDIDO", left=104, top=10, right=170, bottom=28, confidence=82),
        OcrWord(text="62098330", left=260, top=10, right=345, bottom=28, confidence=88),
        OcrWord(text="MODELO", left=10, top=44, right=82, bottom=62, confidence=82),
        OcrWord(text="203302017", left=260, top=44, right=350, bottom=62, confidence=88),
        OcrWord(text="PESO", left=10, top=78, right=58, bottom=96, confidence=82),
        OcrWord(text="BRUTO", left=64, top=78, right=124, bottom=96, confidence=82),
        OcrWord(text="4.3", left=260, top=78, right=292, bottom=96, confidence=88),
        OcrWord(text="KGS", left=300, top=78, right=335, bottom=96, confidence=88),
    ]

    table_text = build_ocr_word_table_text(words)
    fields = {field.key: field.value for field in extract_fields(table_text, source="front_photo")}

    assert fields["po"] == "62098330"
    assert fields["item"] == "203302017"
    assert fields["gw"] == "4.3 KGS"


def test_carton_mark_checks_left_labels_and_right_values_for_side_table():
    from app.schemas.carton_mark import CartonMarkExtractionStatus
    from app.services.carton_mark import compare_label_column, extract_fields

    template_text = (
        "DESCRIPCION VEHICULO DE JUGUETE\n"
        "MODELO 203302017\n"
        "NUMERO DE PEDIDO 62098330\n"
        "PIEZAS POR BULTO 24\n"
        "BULTO 1 DE 1\n"
        "CAJA NUMERO 1 DE 63\n"
        "TALLA NA\n"
        "COLOR MULTICOLOR\n"
        "SECCION / UNECO 424\n"
    )
    photo_text = template_text.replace("NUMERO DE PEDIDO", "NUMERO DE PROVEEDOR")
    status = CartonMarkExtractionStatus(
        source="side_photo",
        ok=True,
        engine="test",
        message="",
        raw_text=photo_text,
    )

    side_fields = {field.key: field.value for field in extract_fields(template_text, source="pdf_side_mark")}
    comparisons = compare_label_column("side", template_text, photo_text, status)
    comparison_map = {item.field_key: item for item in comparisons}

    assert side_fields["carton_no"] == "1 DE 1"
    assert side_fields["box_no"] == "1 DE 63"
    assert side_fields["size"] == "NA"
    assert side_fields["section"] == "424"
    assert comparison_map["left_label:po"].expected == "NUMERO DE PEDIDO"
    assert comparison_map["left_label:po"].actual == "NUMERO DE PROVEEDOR"
    assert comparison_map["left_label:po"].status == "mismatch"
    assert comparison_map["left_label:box_no"].status == "pass"
    assert comparison_map["left_label:section"].status == "pass"


def test_rapidocr_result_recovers_table_fields_from_detected_text_boxes():
    from app.services.carton_mark import extract_fields, rapidocr_result_to_text

    result = SimpleNamespace(
        txts=("NUMERO DE PEDIDO", "62098330", "MODELO", "203302017", "PIEZAS POR BULTO", "24"),
        scores=(0.98, 0.99, 0.97, 0.99, 0.96, 0.99),
        boxes=(
            ((10, 10), (170, 10), (170, 28), (10, 28)),
            ((260, 10), (345, 10), (345, 28), (260, 28)),
            ((10, 44), (82, 44), (82, 62), (10, 62)),
            ((260, 44), (350, 44), (350, 62), (260, 62)),
            ((10, 78), (180, 78), (180, 96), (10, 96)),
            ((260, 78), (292, 78), (292, 96), (260, 96)),
        ),
    )

    text = rapidocr_result_to_text(result)
    fields = {field.key: field.value for field in extract_fields(text, source="rapidocr")}

    assert fields["po"] == "62098330"
    assert fields["item"] == "203302017"
    assert fields["quantity"] == "24"


def test_carton_mark_auto_check_matches_photo_raw_values_when_labels_are_noisy(monkeypatch):
    from app.schemas.carton_mark import CartonMarkExtractionStatus
    from app.services import carton_mark

    def fake_extract_image_text(image_bytes: bytes, *, source: str):
        text = "photo text only 2O 33 20I7 I979I96I9"
        return text, CartonMarkExtractionStatus(
            source=source,
            ok=True,
            engine="test",
            message="",
            raw_text=text,
        )

    monkeypatch.setattr(carton_mark, "extract_image_text", fake_extract_image_text)

    response = carton_mark.build_carton_mark_auto_check(
        pdf_bytes=(
            b"SHIPPING IDENTIFICATION MARK\n"
            b"PO 2033\n"
            b"ITEM 2017\n"
            b"BARCODE 197919619\n"
        ),
        front_image_bytes=b"front-photo",
        side_image_bytes=b"side-photo",
        customer_name="Dickie",
        po="2033",
        item="2017",
    )

    front_statuses = {
        item.field_key: item.status
        for item in response.comparisons
        if item.side == "front"
    }
    assert front_statuses["po"] == "pass"
    assert front_statuses["item"] == "pass"
    assert front_statuses["barcode"] == "pass"


def test_carton_mark_auto_check_backfills_unlabeled_pdf_values_from_photo_fields(monkeypatch):
    from app.schemas.carton_mark import CartonMarkExtractionStatus
    from app.services import carton_mark

    def fake_extract_image_text(image_bytes: bytes, *, source: str):
        text = "COLOR MULTICOLOR\nBARCODE 4197919619"
        return text, CartonMarkExtractionStatus(
            source=source,
            ok=True,
            engine="test",
            message="",
            raw_text=text,
        )

    monkeypatch.setattr(carton_mark, "extract_image_text", fake_extract_image_text)

    response = carton_mark.build_carton_mark_auto_check(
        pdf_bytes=(
            b"SHIPPING IDENTIFICATION MARK\n"
            b"PO 203302017\n"
            b"MULTICOLOR\n"
            b"4197919619\n"
        ),
        front_image_bytes=b"front-photo",
        side_image_bytes=b"side-photo",
        customer_name="Dickie",
        po="203302017",
        item="2017",
    )

    front_expected = {
        field.key: field.value
        for field in response.front_template_fields
    }
    front_statuses = {
        item.field_key: item.status
        for item in response.comparisons
        if item.side == "front"
    }

    assert front_expected["color"] == "MULTICOLOR"
    assert front_statuses["color"] == "pass"


def test_carton_mark_auto_check_tolerates_extracted_identifier_ocr_confusions(monkeypatch):
    from app.schemas.carton_mark import CartonMarkExtractionStatus
    from app.services import carton_mark

    def fake_extract_image_text(image_bytes: bytes, *, source: str):
        text = "PO 2O33\nITEM 20I7\nBARCODE I979I96I9"
        return text, CartonMarkExtractionStatus(
            source=source,
            ok=True,
            engine="test",
            message="",
            raw_text=text,
        )

    monkeypatch.setattr(carton_mark, "extract_image_text", fake_extract_image_text)

    response = carton_mark.build_carton_mark_auto_check(
        pdf_bytes=(
            b"SHIPPING IDENTIFICATION MARK\n"
            b"PO 2033\n"
            b"ITEM 2017\n"
            b"BARCODE 197919619\n"
        ),
        front_image_bytes=b"front-photo",
        side_image_bytes=b"side-photo",
        customer_name="Dickie",
        po="2033",
        item="2017",
    )

    front_statuses = {
        item.field_key: item.status
        for item in response.comparisons
        if item.side == "front"
    }
    assert front_statuses["po"] == "pass"
    assert front_statuses["item"] == "pass"
    assert front_statuses["barcode"] == "pass"


def test_pdf_template_region_selection_keeps_first_front_side_pair():
    from app.services.carton_mark import PdfMarkRegion, select_primary_pdf_mark_regions

    regions = [
        PdfMarkRegion(kind="front", text="FRONT 1", box=(100, 100, 360, 260)),
        PdfMarkRegion(kind="side", text="SIDE 2", box=(380, 100, 520, 260)),
        PdfMarkRegion(kind="front", text="FRONT 3", box=(700, 100, 960, 260)),
        PdfMarkRegion(kind="side", text="SIDE 4", box=(980, 100, 1120, 260)),
    ]

    selected = select_primary_pdf_mark_regions(regions)

    assert [(region.kind, region.text) for region in selected] == [
        ("front", "FRONT 1"),
        ("side", "SIDE 2"),
    ]


def test_pdf_template_side_texts_ignores_repeated_second_pair(monkeypatch):
    from app.schemas.carton_mark import CartonMarkExtractionStatus
    from app.services import carton_mark
    from app.services.carton_mark import PdfMarkRegion

    regions = [
        PdfMarkRegion(kind="front", text="PO 2033\nFRONT 1", box=(100, 100, 360, 260)),
        PdfMarkRegion(kind="side", text="ITEM 2017\nSIDE 2", box=(380, 100, 520, 260)),
        PdfMarkRegion(kind="front", text="PO WRONG\nFRONT 3", box=(700, 100, 960, 260)),
        PdfMarkRegion(kind="side", text="ITEM WRONG\nSIDE 4", box=(980, 100, 1120, 260)),
    ]

    def fake_extract_pdf_mark_regions(pdf_bytes: bytes):
        return regions, CartonMarkExtractionStatus(
            source="pdf_template_regions",
            ok=True,
            engine="test",
            message="",
            raw_text="",
        )

    monkeypatch.setattr(carton_mark, "extract_pdf_mark_regions", fake_extract_pdf_mark_regions)

    front_text, side_text, status = carton_mark.extract_pdf_template_side_texts(
        b"pdf",
        fallback_text="fallback",
    )

    assert status.ok
    assert "FRONT 1" in front_text
    assert "FRONT 3" not in front_text
    assert "SIDE 2" in side_text
    assert "SIDE 4" not in side_text


def test_pdf_vector_regions_keep_only_the_leftmost_front_and_side_marks():
    from app.services.carton_mark import (
        PdfTextSpan,
        build_pdf_vector_mark_regions,
        classify_pdf_mark_regions,
        extract_fields,
        select_primary_pdf_mark_regions,
    )

    spans = [
        PdfTextSpan(text="SHIPPING IDENTIFICATION MARK", x=100, y=340),
        PdfTextSpan(text="IMPORTADOR DILISA", x=100, y=315),
        PdfTextSpan(text="62098330NUMERO DE PEDIDO", x=100, y=290),
        PdfTextSpan(text="MODELO 203302017", x=100, y=265),
        PdfTextSpan(text="DESCRIPCION", x=500, y=340),
        PdfTextSpan(text="VEHICULO DE JUGUETE", x=650, y=340),
        PdfTextSpan(text="NUMERO DE PEDIDO", x=500, y=315),
        PdfTextSpan(text="62098330", x=650, y=315),
        PdfTextSpan(text="PIEZAS POR BULTO", x=500, y=290),
        PdfTextSpan(text="24 24", x=650, y=290),
        PdfTextSpan(text="BULTO", x=500, y=265),
        PdfTextSpan(text="1 DE 1", x=650, y=265),
        PdfTextSpan(text="SHIPPING IDENTIFICATION MARK", x=900, y=340),
        PdfTextSpan(text="IMPORTADOR OTRO CLIENTE", x=900, y=315),
        PdfTextSpan(text="99999999NUMERO DE PEDIDO", x=900, y=290),
        PdfTextSpan(text="DESCRIPCION", x=1300, y=340),
        PdfTextSpan(text="OTRO PRODUCTO", x=1450, y=340),
        PdfTextSpan(text="NUMERO DE PEDIDO", x=1300, y=315),
        PdfTextSpan(text="99999999", x=1450, y=315),
        PdfTextSpan(text="PIEZAS POR BULTO", x=1300, y=290),
        PdfTextSpan(text="12", x=1450, y=290),
    ]

    regions = classify_pdf_mark_regions(
        build_pdf_vector_mark_regions(spans, page_width=1700, page_height=400),
    )
    selected = select_primary_pdf_mark_regions(regions)
    front_fields = {field.key: field.value for field in extract_fields(selected[0].text, source="pdf_front_mark")}
    side_fields = {field.key: field.value for field in extract_fields(selected[1].text, source="pdf_side_mark")}

    assert [(region.kind, region.box[0]) for region in selected] == [("front", 100), ("side", 500)]
    assert front_fields["po"] == "62098330"
    assert front_fields["item"] == "203302017"
    assert side_fields["po"] == "62098330"
    assert side_fields["quantity"] == "24"
    assert side_fields["carton_no"] == "1 DE 1"


def _install_multi_page_photo_check_fakes(monkeypatch, carton_mark, photo_text: str):
    from app.schemas.carton_mark import CartonMarkExtractionStatus
    from app.services.carton_mark import PdfMarkRegion

    def status(source: str) -> CartonMarkExtractionStatus:
        return CartonMarkExtractionStatus(source=source, ok=True, engine="test", raw_text="")

    regions = [
        PdfMarkRegion(kind="front", text="PO 4500\nITEM 1001", box=(10, 10, 300, 200), page_index=0),
        PdfMarkRegion(kind="side", text="ITEM 1001\nG.W. 4 KGS\nN.W. 3 KGS", box=(320, 10, 500, 200), page_index=0),
        PdfMarkRegion(kind="front", text="PO 4500\nITEM 2002", box=(10, 10, 300, 200), page_index=1),
        PdfMarkRegion(kind="side", text="ITEM 2002\nG.W. 5 KGS\nN.W. 4 KGS", box=(320, 10, 500, 200), page_index=1),
    ]
    monkeypatch.setattr(carton_mark, "extract_pdf_text", lambda _content: (
        "\n".join(region.text for region in regions),
        status("pdf_template"),
    ))
    monkeypatch.setattr(carton_mark, "extract_pdf_mark_regions", lambda _content: (
        regions,
        status("pdf_template_regions"),
    ))
    monkeypatch.setattr(carton_mark, "extract_image_text", lambda _content, *, source: (
        photo_text,
        status(source),
    ))


def test_batch_photo_anchor_selects_second_pdf_page(monkeypatch):
    from app.services import carton_mark

    _install_multi_page_photo_check_fakes(monkeypatch, carton_mark, "ITEM 2002")
    response = carton_mark.build_carton_mark_batch_auto_check(
        pdf_bytes=b"pdf",
        front_images=[("carton.jpg", b"photo")],
        side_images=[],
    )

    result = response.items[0].result
    assert {field.key: field.value for field in result.template_fields}["item"] == "2002"
    selection = next(status for status in result.extraction if "photo-anchor" in status.engine)
    assert "第 2 页" in selection.message
    assert "ITEM=2002" in selection.message
    assert selection.matched_page == 2
    assert selection.page_count == 2
    assert selection.match_confidence is not None and selection.match_confidence >= 0.9
    assert selection.requires_review is False
    assert result.summary.overall_status == "需复核"
    assert any(item.field_key == "photo_coverage" for item in result.comparisons)


def test_batch_multi_page_without_unique_anchor_requires_review(monkeypatch):
    from app.services import carton_mark

    _install_multi_page_photo_check_fakes(monkeypatch, carton_mark, "G.W. 9 KGS")
    response = carton_mark.build_carton_mark_batch_auto_check(
        pdf_bytes=b"pdf",
        front_images=[("carton.jpg", b"photo")],
        side_images=[],
    )

    result = response.items[0].result
    assert result.summary.overall_status == "需复核"
    assert result.summary.mismatch_count == 0
    assert result.summary.review_count > 0
    selection = next(status for status in result.extraction if status.engine == "pdf-page-selection")
    assert not selection.ok
    assert "未自动选择页面" in selection.message
    assert selection.matched_page is None
    assert selection.page_count == 2
    assert selection.requires_review is True
    assert selection.review_reason == selection.message


def test_front_upload_can_compare_observed_side_weight_on_selected_page(monkeypatch):
    from app.services import carton_mark

    _install_multi_page_photo_check_fakes(monkeypatch, carton_mark, "ITEM 2002\nG.W. 9 KGS")
    response = carton_mark.build_carton_mark_batch_auto_check(
        pdf_bytes=b"pdf",
        front_images=[("unfolded-carton.jpg", b"photo")],
        side_images=[],
    )

    comparison = next(
        item for item in response.items[0].result.comparisons if item.field_key == "gw"
    )
    assert comparison.side == "side"
    assert comparison.expected == "5 KGS"
    assert comparison.actual == "9 KGS"
    assert comparison.status == "mismatch"


def test_single_face_photo_does_not_report_unseen_other_face_fields(monkeypatch):
    from app.services import carton_mark

    _install_multi_page_photo_check_fakes(monkeypatch, carton_mark, "ITEM 2002")
    response = carton_mark.build_carton_mark_batch_auto_check(
        pdf_bytes=b"pdf",
        front_images=[("front-only.jpg", b"photo")],
        side_images=[],
    )

    comparisons = response.items[0].result.comparisons
    assert not any(item.field_key in {"gw", "nw"} for item in comparisons)
    assert not any(item.status == "missing_actual" for item in comparisons)


def test_carton_certification_and_pi_fields_do_not_get_swallowed_by_gw():
    from app.services.carton_mark import extract_fields

    fields = {
        field.key: field.value
        for field in extract_fields(
            (
                "PI NO. 1100036756\n"
                "DOUBLE WALL\n"
                "BURSTING TEST 200 LBS PER SQ IN\n"
                "MIN COMB WT FACINGS 92 LBS PER M SQ FT\n"
                "SIZE LIMIT 75 INCHES\n"
                "GROSS WT LT 65 LBS\n"
            ),
            source="photo",
        )
    }

    assert fields["pi_no"] == "1100036756"
    assert fields["wall_type"] == "DOUBLE WALL"
    assert fields["bursting_test"] == "200"
    assert fields["min_combined_weight"] == "92"
    assert fields["size_limit"] == "75"
    assert fields["gross_weight_limit"] == "65"
    assert "gw" not in fields


def test_pdf_document_line_filter_keeps_two_digit_carton_totals():
    from app.services.carton_mark import _clean_pdf_document_line

    assert _clean_pdf_document_line("39") == "39"
    assert _clean_pdf_document_line("49") == "49"
    assert _clean_pdf_document_line("PAP") == ""
    assert _clean_pdf_document_line("RESY") == ""


def test_trailing_item_label_selects_exact_adjacent_pdf_page():
    from app.schemas.carton_mark import CartonMarkExtractionStatus
    from app.services.carton_mark import PdfMarkRegion, select_pdf_mark_page

    regions = [
        PdfMarkRegion(kind="full", text="ARTICLE NO.: 100371\nNUMBER OF PIECES: 6 PCS", box=(0, 0, 800, 600), page_index=0),
        PdfMarkRegion(kind="full", text="ARTICLE NO.: 100372\nNUMBER OF PIECES: 6 PCS", box=(0, 0, 800, 600), page_index=1),
        PdfMarkRegion(kind="full", text="ARTICLE NO.: 100373\nNUMBER OF PIECES: 6 PCS", box=(0, 0, 800, 600), page_index=2),
    ]
    status = CartonMarkExtractionStatus(source="pdf", ok=True, engine="test", raw_text="")

    selected, selection = select_pdf_mark_page(
        regions,
        photo_text="ARTICLE NO.: 100372 NUMBER OF PIECES: 6 PCS MADE IN CHINA",
        metadata={},
        extraction_status=status,
    )

    assert selection.matched_page == 2
    assert all(region.page_index == 1 for region in selected)
    assert "ITEM=100372" in selection.message


def test_pdf_layout_text_recovers_pi_weights_and_measurement():
    from app.services.carton_mark import extract_fields, normalize_pdf_layout_text

    layout_text = (
        "Carton      Net      W eight:5.04 KGS\n"
        "Carton      G  ross      W eight:5.93 KGS\n"
        "Carton      S  ize(cm):56.5 x 31.4 x 63.5 CM\n"
        "PINo.      1100036670\n"
        "Qty/Ctn:      6 PCS\n"
    )
    fields = {
        field.key: field.value
        for field in extract_fields(normalize_pdf_layout_text(layout_text), source="pdf")
    }

    assert fields["pi_no"] == "1100036670"
    assert fields["gw"] == "5.93 KGS"
    assert fields["nw"] == "5.04 KGS"
    assert fields["measurement"] == "56.5 x 31.4 x 63.5 CM"
    assert fields["quantity"] == "6"


def test_pdf_layout_text_repairs_labels_that_touch_their_values():
    from app.services.carton_mark import (
        extract_field_labels,
        extract_fields,
        normalize_pdf_layout_text,
    )

    normalized = normalize_pdf_layout_text(
        "DESCRIPTIONRB AXOLOTLCORNS SURP D: 45.10 X 23.20 X 40.30 CMS\n"
        "*Keycode 43569786 GRS WT3.18 KGS\n"
        "CARTON NO.of 5\n"
    )
    fields = {
        field.key: field.value
        for field in extract_fields(normalized, source="pdf")
    }
    labels = extract_field_labels(normalized)

    assert fields["description"].startswith("RB AXOLOTLCORNS SURP")
    assert fields["gw"] == "3.18 KGS"
    assert fields["measurement"] == "45.10 X 23.20 X 40.30 CMS"
    assert labels["carton_no"].upper() == "CARTON NO. OF"
    assert "measurement" not in labels


def test_equivalent_field_label_aliases_do_not_create_false_mismatches():
    from app.services.carton_mark import field_label_values_match

    assert field_label_values_match("ITEM", "Item No.")
    assert field_label_values_match("Carton No. of", "Carton No.")
    assert not field_label_values_match("PO NUMBER", "ARTICLE NUMBER")


def test_certification_stamp_accepts_wall_type_without_ocr_space():
    from app.services.carton_mark import extract_carton_certification_values

    values = extract_carton_certification_values(
        "DOUBLEWALL BOX MEET ALL CONSTRUCTION REQUIREMENTS "
        "BURSTING TEST 200 MIN COMB WT FACINGS 92 SIZE LIMIT 75 GROSS WT LT 65"
    )

    assert values == {
        "wall_type": "DOUBLE WALL",
        "bursting_test": "200",
        "min_combined_weight": "92",
        "size_limit": "75",
        "gross_weight_limit": "65",
    }


def test_carton_number_of_layout_is_equivalent_but_missing_item_needs_review(monkeypatch):
    from app.schemas.carton_mark import CartonMarkExtractionStatus
    from app.services import carton_mark
    from app.services.carton_mark import PdfMarkRegion

    status = CartonMarkExtractionStatus(source="test", ok=True, engine="test", raw_text="")
    monkeypatch.setattr(carton_mark, "extract_pdf_text", lambda _content: (
        "Order No. 21894319\nDescription RB AXOLOTLCORNS SURP\nCarton No. of 5",
        status,
    ))
    monkeypatch.setattr(carton_mark, "extract_pdf_mark_regions", lambda _content: ([
        PdfMarkRegion(
            kind="full",
            text="Order No. 21894319\nDescription RB AXOLOTLCORNS SURP\nCarton No. of 5",
            box=(0, 0, 800, 600),
        ),
    ], status))
    monkeypatch.setattr(carton_mark, "extract_image_text", lambda _content, *, source: (
        "Order No. 21894319\nDescription RB AXOLOTLCORNS SURP\nCarton No. 5",
        status.model_copy(update={"source": source}),
    ))

    response = carton_mark.build_carton_mark_batch_auto_check(
        pdf_bytes=b"pdf",
        front_images=[("photo.jpg", b"photo")],
        side_images=[],
    )
    comparisons = response.items[0].result.comparisons

    carton = next(item for item in comparisons if item.field_key == "carton_no")
    assert carton.status == "pass"
    assert not any(item.field_key == "left_label:carton_no" and item.status == "mismatch" for item in comparisons)
    assert any(item.field_key == "photo_coverage" and item.status == "review" for item in comparisons)
    assert response.summary.overall_status == "需复核"


def test_unlabelled_order_and_contract_numbers_are_not_barcodes():
    from app.services.carton_mark import extract_fields

    fields = {
        field.key: field.value
        for field in extract_fields(
            "Order No. 4500211063\nContract Number: 21216364\nITEM 100372",
            source="photo",
        )
    }

    assert fields["item"] == "100372"
    assert "barcode" not in fields


def test_single_page_wrong_item_remains_anomaly(monkeypatch):
    from app.schemas.carton_mark import CartonMarkExtractionStatus
    from app.services import carton_mark
    from app.services.carton_mark import PdfMarkRegion

    status = CartonMarkExtractionStatus(source="test", ok=True, engine="test", raw_text="")
    template = "PO NUMBER 484860\nITEM NO. 415647\nDESCRIPTION FUGGLER BADDIE"
    monkeypatch.setattr(carton_mark, "extract_pdf_text", lambda _content: (template, status))
    monkeypatch.setattr(carton_mark, "extract_pdf_mark_regions", lambda _content: ([
        PdfMarkRegion(kind="full", text=template, box=(0, 0, 800, 600)),
    ], status))
    monkeypatch.setattr(carton_mark, "extract_image_text", lambda _content, *, source: (
        "ARTICLE NUMBER 7496.100.000.00\nG.W. 3.01 KGS\nN.W. 2.48 KGS",
        status.model_copy(update={"source": source}),
    ))

    response = carton_mark.build_carton_mark_batch_auto_check(
        pdf_bytes=b"pdf",
        front_images=[("wrong-set.jpg", b"photo")],
        side_images=[],
    )

    item = next(item for item in response.items[0].result.comparisons if item.field_key == "item")
    assert item.status == "mismatch"
    assert response.summary.overall_status == "发现异常"


def test_structured_certification_values_suppress_duplicate_label_ocr_noise():
    from app.schemas.carton_mark import CartonMarkExtractionStatus
    from app.services.carton_mark import compare_observed_photo_fields, extract_fields

    template_text = (
        "WALL TYPE SINGLE WALL\n"
        "BURSTING TEST 175\n"
        "MIN COMB WT FACINGS 75\n"
        "SIZE LIMIT 60\n"
        "GROSS WT LT 40\n"
    )
    photo_text = (
        "BOARD TYPE DOUBLE WALL\n"
        "BUASTING TEST 200\n"
        "MIN COMBINED WEIGHT FACINGS 92\n"
        "SIZE LT 75\n"
        "GROSS WEIGHT LIMIT 65\n"
    )
    status = CartonMarkExtractionStatus(
        source="front_photo",
        ok=True,
        engine="test",
        raw_text=photo_text,
    )
    comparisons = compare_observed_photo_fields(
        "front",
        photo_text=photo_text,
        photo_status=status,
        photo_fields=extract_fields(photo_text, source="front_photo"),
        front_template_text="",
        side_template_text=template_text,
        front_template_fields=[],
        side_template_fields=extract_fields(template_text, source="pdf_side_mark"),
    )

    value_mismatches = {
        item.field_key
        for item in comparisons
        if item.comparison_scope == "right_value" and item.status == "mismatch"
    }
    assert value_mismatches == {
        "wall_type",
        "bursting_test",
        "min_combined_weight",
        "size_limit",
        "gross_weight_limit",
    }
    assert not any(
        item.field_key.startswith("left_label:")
        and item.field_key.removeprefix("left_label:") in value_mismatches
        for item in comparisons
    )


def test_repeated_photo_panels_report_one_missing_country_origin_face():
    from app.schemas.carton_mark import CartonMarkExtractionStatus
    from app.services.carton_mark import (
        compare_observed_photo_fields,
        extract_fields,
        rapidocr_photo_panel_evidence_lines,
    )

    def box(left: int, top: int, right: int, bottom: int):
        return [[left, top], [right, top], [right, bottom], [left, bottom]]

    result = SimpleNamespace(
        txts=("ITEM NO.: 15755", "ITEM NO.: 15755", "MADE IN CHINA"),
        boxes=(box(10, 10, 150, 40), box(10, 210, 150, 240), box(10, 250, 170, 280)),
        scores=(0.99, 0.99, 0.99),
    )
    evidence = rapidocr_photo_panel_evidence_lines(result)
    assert "[PHOTO_PANEL_COUNT:2]" in evidence
    assert "[PHOTO_COUNTRY_PANEL:1:CHINA]" in evidence

    photo_text = "\n".join(["ITEM NO.: 15755", *evidence])
    template_text = "ITEM NO.: 15755\nMADE IN CHINA"
    status = CartonMarkExtractionStatus(
        source="front_photo",
        ok=True,
        engine="test",
        raw_text=photo_text,
    )
    comparisons = compare_observed_photo_fields(
        "front",
        photo_text=photo_text,
        photo_status=status,
        photo_fields=extract_fields(photo_text, source="front_photo"),
        front_template_text=template_text,
        side_template_text=template_text,
        front_template_fields=extract_fields(template_text, source="pdf_front_mark"),
        side_template_fields=extract_fields(template_text, source="pdf_side_mark"),
    )

    origin = next(item for item in comparisons if item.field_key == "country_of_origin_panels")
    assert origin.status == "missing_actual"
    assert origin.expected == "CHINA / CHINA（2 面）"
    assert origin.actual == "CHINA（1 面）"
    assert "1/2 面" in origin.note
