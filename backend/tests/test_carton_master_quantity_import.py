"""Packaging template rules share formal-order calculation and retain old files."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from decimal import Decimal as D
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from test_molding_sample_api import make_client, login_as
from test_carton_master_api import prepare, read
from test_carton_master_import_api import upload
from test_carton_procurement_api import _order_payload

BASE = '/api/carton-procurement'
LEGACY14 = ['货号', '产品名称', '配置组', '纸品类型', '纸质', '长', '宽', '高',
            '尺寸单位', '计量单位', '每箱装产品数', '推荐', '状态', '备注']
LEGACY16 = [*LEGACY14, '每箱净重kg', '每箱毛重kg']
CURRENT20 = [*LEGACY16, '纸品行号', '计算方式', '对应箱行号', '每箱配套用量']


@pytest.fixture(autouse=True)
def isolated(monkeypatch, tmp_path):
    import test_molding_sample_api
    monkeypatch.setattr(test_molding_sample_api, 'TEST_TMP_DIR', tmp_path)


def paper(*, code='00123', group='标准装', no=1, name='外箱', packing=24,
          mode='按装箱数', parent=None, multiple=None, height=15, unit='个',
          net=None, gross=None):
    return [code, '产品', group, name, 'A33', 30, 20, height, 'cm', unit, packing,
            '否', '启用', '完整配置示例', net, gross, no, mode, parent, multiple]


def card(**values):
    return paper(no=2, name='平卡', packing=None, mode='按箱配套', parent=1,
                 multiple=2, height=None, unit='片', **values)


def book_bytes(rows, headers=CURRENT20):
    book = Workbook()
    book.active.title = '导入数据'
    book.active.append(headers)
    for row in rows:
        book.active.append(row)
    stream = BytesIO(); book.save(stream)
    return stream.getvalue()


def planned(tmp_path, rows, *, kind='configurations', headers=CURRENT20):
    from sqlalchemy import create_engine, select
    from sqlalchemy.orm import Session
    from app.models.carton_master import CartonMasterRecord
    from app.services import carton_master_import as service
    engine = create_engine('sqlite:///' + (tmp_path / 'master-plan.db').as_posix())
    CartonMasterRecord.__table__.create(engine)
    try:
        entries, errors = service.parse(book_bytes(rows, headers), '包装.xlsx', kind)
        with Session(engine) as db:
            result = service.plan(db, 'huaxing', kind, entries, errors)
            assert db.scalar(select(CartonMasterRecord.id)) is None
            assert not db.new and not db.dirty
            return result
    finally:
        engine.dispose()


@pytest.mark.parametrize('change, expected', [
    ('missing-number', '行号'), ('duplicate-number', '行号'), ('gap', '行号'),
    ('self', '关联'), ('out-of-range', '关联'), ('card-chain', '关联'),
    ('mixed-packing', '留空'), ('packing-with-parent', '留空'),
    ('zero-multiple', '正数'), ('negative-multiple', '正数'),
    ('precision', '4 位'), ('too-large', '14 位'), ('nan', '正数'),
    ('fraction-number', '整数'), ('formula', '公式'),
    ('missing-width', '全部填写'), ('zero-height', '正数'),
    ('bad-parent-row', '毛重'), ('cross-group', '行号'), ('unknown-mode', '计算方式'),
])
def test_preview_rejects_invalid_group_graphs_and_rows_before_any_write(tmp_path, change, expected):
    rows = [paper(), card()]
    if change == 'missing-number': rows[0][16] = None
    elif change == 'duplicate-number': rows[1][16] = 1
    elif change == 'gap': rows[1][16] = 3
    elif change == 'self': rows[1][18] = 2
    elif change == 'out-of-range': rows[1][18] = 3
    elif change == 'card-chain': rows[0][10], rows[0][17:20] = None, ['按箱配套', 2, 1]
    elif change == 'mixed-packing': rows[1][10] = 24
    elif change == 'packing-with-parent': rows[0][18] = 2
    elif change == 'zero-multiple': rows[1][19] = 0
    elif change == 'negative-multiple': rows[1][19] = -1
    elif change == 'precision': rows[1][19] = '0.12345'
    elif change == 'too-large': rows[1][19] = '100000000000000'
    elif change == 'nan': rows[1][19] = 'NaN'
    elif change == 'fraction-number': rows[1][16] = '1.5'
    elif change == 'formula': rows[1][18] = '=1'
    elif change == 'missing-width': rows[0][6] = None
    elif change == 'zero-height': rows[0][7] = 0
    elif change == 'bad-parent-row': rows[0][14:16] = [3, 2]
    elif change == 'cross-group': rows[1][2] = '另一个组'
    elif change == 'unknown-mode': rows[1][17] = '本次手动数量'
    result, _ = planned(tmp_path, rows)
    assert expected in '；'.join(result['errors']), result


def test_scalar_options_accept_two_or_three_dimensions_and_keep_blank_height(tmp_path):
    result, writes = planned(tmp_path, [['平卡', 'A33', 30, 20, None],
        ['外箱', 'B33', 30, 20, 15]], kind='paper-options',
        headers=['纸品类型', '纸质', '长', '宽', '高'])
    assert not result['errors']
    assert writes[0][1].data.specifications == ['30*20', '30*20*15']


def test_download_reverse_interleaved_groups_apply_and_reuse_actual_order_calculation(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        response = client.get(BASE + '/master-data/import/configurations/template', params={'factory_id': 'huaxing'})
        assert response.status_code == 200, response.text
        book = load_workbook(BytesIO(response.content))
        assert [cell.value for cell in book['导入数据'][1]] == CURRENT20
        assert book['导入数据'].max_row == 1
        rows = [card(), paper(code='00024', no=2, name='内箱', packing=12),
                paper(code='00024', no=3, name='刀卡', packing=None, mode='按箱配套',
                    parent=2, multiple=3, height=None, unit='片'),
                paper(net='0', gross='1.2'), paper(code='00024', no=1),
                paper(code='00999', no=2, name='A型箱', packing=12),
                paper(code='00999', no=1, name='Z型箱', packing=24)]
        for row in rows: book['导入数据'].append(row)
        stream = BytesIO(); book.save(stream); content = stream.getvalue()
        preview = upload(client, 'configurations', content)
        assert preview.status_code == 200, preview.text
        preview = preview.json()
        assert preview['errors'] == [] and preview['added'] == 3
        assert any('第2行配套第1行，每箱2片' in value for value in preview['details'])
        assert any('第3行配套第2行，每箱3片' in value for value in preview['details'])
        assert read(client)['records'] == []
        applied = upload(client, 'configurations', content, preview['preview_token'])
        assert applied.status_code == 200, applied.text
        configs = {row['code']: row for row in read(client)['records'] if row['kind'] == 'CONFIG'}
        assert len(configs) == 3
        plain_lines = configs['00999']['data']['lines']
        assert [line['packaging_type'] for line in plain_lines] == ['A型箱', 'Z型箱']
        for number, line in enumerate(plain_lines, 1):
            assert any(f"00999 / 标准装：第{number}行 {line['packaging_type']}" in value for value in preview['details'])
        from app.services.carton_quantity import requirements
        assert requirements(100, configs['00024']['data']['lines']) == [5, 9, 27]
        selected = configs['00123']
        assert [row['specification'] for row in selected['data']['lines']] == ['30*20*15', '30*20']
        assert D(selected['data']['lines'][0]['net_weight_kg']) == 0
        assert D(selected['data']['lines'][0]['gross_weight_kg']) == D('1.2')
        assert client.get(BASE + '/inventory/movements', params={'factory_id': 'huaxing'}).json()['items'] == []
        data = {**_order_payload(), 'item_no': '00123', 'product_name': '产品',
            'product_order_quantity': '100', 'master_config_id': selected['id'],
            'master_config_revision': selected['revision'],
            'lines': [{**row, 'unit_price': '2'} for row in selected['data']['lines']]}
        made = client.post(BASE + '/orders', json=data)
        assert made.status_code == 201, made.text
        assert [D(row['required_quantity']) for row in made.json()['lines']] == [5, 10]
        assert made.json()['lines'][1]['specification'] == '30*20'
        # Existing manually maintained equivalent decimal scales also skip;
        # import comparison must not rewrite that record or its old identity.
        def maintain(row, data, reason):
            response = client.patch(BASE + f"/master-data/{row['id']}", json=dict(
                factory_id='huaxing', kind='CONFIG', code=row['code'], data=data,
                expected_revision=row['revision'], status=row['status'], preferred=row['preferred'], reason=reason))
            assert response.status_code == 200, response.text
            return response.json()
        edited_data = deepcopy(selected['data'])
        edited_data['lines'][0]['quantity_rule'] = dict(mode='PACKING', manual_quantity=None)
        edited_data['lines'][1]['quantity_rule']['per_carton_quantity'] = '2.0000'
        edited = maintain(selected, edited_data, '已有手工配置使用等价四位小数及显式父箱规则')
        assert requirements(100, edited['data']['lines']) == [5, 10]
        plain = configs['00999']
        manual = deepcopy(plain['data'])
        manual['lines'][0].update(usage_quantity=None, quantity_rule=dict(
            mode='PER_CARTON', carton_line_no=2, per_carton_quantity='1'))
        switched = maintain(plain, manual, '手工将第一纸品切换为按箱配套')
        manual = deepcopy(switched['data'])
        manual['lines'][0].update(usage_quantity='12', quantity_rule=dict(mode='PACKING', manual_quantity=None))
        switched_back = maintain(switched, manual, '手工切回装箱数保留显式规则')
        assert switched_back['data']['lines'][0]['quantity_rule']['mode'] == 'PACKING'
        before = read(client)['records']
        equivalent_rows = deepcopy(rows)
        equivalent_rows[0][19] = '2.0000'
        equivalent_rows[2][19] = '3.0000'
        equivalent = book_bytes(equivalent_rows)
        duplicate = upload(client, 'configurations', equivalent).json()
        assert duplicate['errors'] == [] and duplicate['added'] == 0 and duplicate['skipped'] == 3
        assert upload(client, 'configurations', equivalent, duplicate['preview_token']).status_code == 200
        assert read(client)['records'] == before


def test_exact_legacy_14_and_16_headers_reuse_new_packing_semantics_without_overwriting_master(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        for width, headers in ((14, LEGACY14), (16, LEGACY16)):
            code = '00' + str(width)
            row = paper(code=code, no=None, mode='', net='1' if width == 16 else None,
                        gross='2' if width == 16 else None)
            content = book_bytes([row[:width]], headers)
            preview = upload(client, 'configurations', content).json()
            assert preview['errors'] == [] and preview['added'] == 1
            applied = upload(client, 'configurations', content, preview['preview_token'])
            assert applied.status_code == 200, applied.text
            before = read(client)['records']
            for explicit_mode in ('', '按装箱数'):
                current_row = deepcopy(row); current_row[17] = explicit_mode
                current_row[11:14] = ['是', '停用', '新文件不能覆盖已维护资料']
                current_content = book_bytes([current_row])
                repeated = upload(client, 'configurations', current_content).json()
                assert repeated['errors'] == [] and repeated['added'] == 0 and repeated['skipped'] == 1
                assert upload(client, 'configurations', current_content, repeated['preview_token']).status_code == 200
                assert read(client)['records'] == before
        assert len(read(client)['records']) == 2


def test_bad_parent_cannot_renumber_child_and_valid_group_cannot_partially_apply(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        invalid = [paper(net=2, gross=1), card(), paper(code='HEALTHY', no=None)]
        content = book_bytes(invalid)
        preview = upload(client, 'configurations', content).json()
        assert any('毛重' in error for error in preview['errors'])
        assert any('行号' in error for error in preview['errors'])
        assert upload(client, 'configurations', content, preview['preview_token']).status_code == 422
        assert read(client)['records'] == []
        fixed = book_bytes([paper(), card()])
        preview = upload(client, 'configurations', fixed).json()
        assert preview['errors'] == []
        from sqlalchemy import select
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonAuditEvent
        def audits():
            with SessionLocal() as db:
                return {(event.id, event.event_type) for event in db.scalars(select(CartonAuditEvent))}
        before_audits = audits()
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(lambda _: upload(client, 'configurations', fixed,
                preview['preview_token']), range(2)))
        assert sorted(response.status_code for response in responses) in ([200, 409], [200, 429])
        # The bounded parser can reject before the factory transaction; once it
        # is free the original successful preview is still stale and cannot run.
        assert upload(client, 'configurations', fixed, preview['preview_token']).status_code == 409
        assert len(read(client)['records']) == 1
        assert client.get(BASE + '/inventory/movements', params={'factory_id': 'huaxing'}).json()['items'] == []
        added_audits = audits() - before_audits
        assert len(added_audits) == 1 and next(iter(added_audits))[1] == 'MASTER_DATA_SAVED'


def test_apply_rereads_master_permission_after_factory_lock_without_record_or_audit_write(monkeypatch):
    monkeypatch.setenv('AUTHZ_MODE', 'enforce')
    with make_client(monkeypatch) as client:
        prepare(client)
        profile = login_as(client, 'warehouse_keeper')
        content = book_bytes([paper(), card()])
        preview = upload(client, 'configurations', content)
        assert preview.status_code == 200, preview.text
        assert not preview.json()['errors']
        from sqlalchemy import select
        from app.db import SessionLocal
        from app.models.auth import AuthPermission, AuthUserPermissionOverride, AuthUserAuthorizationRevision
        from app.models.carton_master import CartonMasterRecord
        from app.models.carton_procurement import CartonAuditEvent
        from app.services import carton_procurement as core
        def snapshot():
            with SessionLocal() as db:
                return ([row.id for row in db.scalars(select(CartonMasterRecord).order_by(CartonMasterRecord.id))],
                        [row.id for row in db.scalars(select(CartonAuditEvent).order_by(CartonAuditEvent.id))])
        before = snapshot()
        original = core._lock_receipt_factory
        touched = []
        def revoke_after_lock(db, factory):
            original(db, factory)
            permission = db.scalar(select(AuthPermission).where(AuthPermission.code == 'carton_procurement:master_manage'))
            db.add(AuthUserPermissionOverride(id='master-import-lock-deny', user_id=profile['id'],
                permission_id=permission.id, effect='deny', factory_id='huaxing', department='*',
                status='active', reason='模拟等待厂区锁期间撤回维护权限'))
            revision = db.get(AuthUserAuthorizationRevision, profile['id'])
            if revision is None:
                db.add(AuthUserAuthorizationRevision(user_id=profile['id'], revision=2, updated_at='2026-10-20'))
            else:
                revision.revision += 1
            db.flush(); touched.append(True)
        monkeypatch.setattr(core, '_lock_receipt_factory', revoke_after_lock)
        denied = upload(client, 'configurations', content, preview.json()['preview_token'])
        assert denied.status_code == 403, denied.text
        assert touched
        monkeypatch.setattr(core, '_lock_receipt_factory', original)
        assert snapshot() == before
