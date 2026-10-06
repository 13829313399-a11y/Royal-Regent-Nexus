"""A late vendor document adds evidence, never a second physical receipt/payable."""
import json
from decimal import Decimal
from uuid import uuid4

import pytest
from sqlalchemy import select

from test_carton_supplier_portal import BASE, setup_portal, accept_all, supplier_login, ship_payload, receive_payload, _dongkang_delivery_file
from test_molding_sample_api import make_client, login_as


def manual_receipt(client, order, *, note="MANUAL-01", rejected_last=False, full=False):
    login_as(client, "admin")
    bins = client.get("/api/carton-procurement/inventory/locations", params={"factory_id": "huaxing"}).json()
    location = next(row["id"] for row in bins if row["bin_code"] == "A-01")
    quantities = [float(line["required_quantity"]) if full else 10 for line in order["lines"]]
    lines = [{"order_line_id": line["id"], "delivered_quantity": quantities[index], "received_quantity": quantities[index],
        "unit_price": 2, "rejected_quantity": 10 if rejected_last and index == len(order["lines"])-1 else 0,
        "location_allocations": [] if rejected_last and index == len(order["lines"])-1 else [{"location_id": location, "quantity": quantities[index]}]}
        for index, line in enumerate(order["lines"])]
    response = client.post("/api/carton-procurement/receipts", json={"factory_id": "huaxing", "request_id": str(uuid4()),
        "post_immediately": True, "delivery_note_no": note, "delivery_date": "2026-09-21", "acceptance_date": "2026-09-21", "lines": lines})
    assert response.status_code == 201, response.text
    return response.json()


def late_note(client, order, *, mode="EXISTING_RECEIPT", content=None):
    supplier_login(client)
    content = content or _dongkang_delivery_file(order)
    preview = client.post(BASE + "/shipments/import-preview", files={"file": ("dongkang.xlsx", content)}).json()
    assert preview["groups"][0]["receipt_link_ready"], preview
    data = {"sha256": preview["sha256"], "selections": json.dumps([{"factory_id": "huaxing", "delivery_note_no": "DK-IMPORT-01", "registration_mode": mode}])}
    response = client.post(BASE + "/shipments/import-confirm", files={"file": ("dongkang.xlsx", content)}, data=data)
    assert response.status_code == 200, response.text
    return response.json()["shipments"][0], content, data


def link_body(shipment, receipt):
    return {"factory_id": "huaxing", "request_id": str(uuid4()), "expected_revision": shipment["revision"],
        "receipt_id": receipt["id"], "expected_receipt_revision": receipt["revision"], "reason": "已手动入库，供应商后补原凭证"}


def evidence():
    from app.db import SessionLocal
    from app.models.carton_procurement import CartonReceipt, CartonInventoryMovement
    from app.services.carton_supplier_settlement import sources
    with SessionLocal() as db:
        receipts = list(db.scalars(select(CartonReceipt).where(CartonReceipt.factory_id == "huaxing")))
        movements = [(row.id, str(row.quantity), str(row.unit_price), row.occurred_at) for row in db.scalars(select(CartonInventoryMovement).where(CartonInventoryMovement.factory_id == "huaxing").order_by(CartonInventoryMovement.id))]
        rows, _, _, _ = sources(db, "huaxing", receipts[0].supplier_id, "2026-09", "CNY")
        return len(receipts), movements, rows


@pytest.mark.parametrize("same_number,rejected_last", [(False, False), (True, False), (False, True)])
def test_late_link_keeps_original_stock_payable_and_acceptance(monkeypatch, same_number, rejected_last):
    with make_client(monkeypatch) as client:
        setup_portal(client); order = accept_all(client)
        receipt = manual_receipt(client, order, note="DK-IMPORT-01" if same_number else "MANUAL-01", rejected_last=rejected_last)
        before = evidence()
        shipment, content, data = late_note(client, order)
        assert shipment["requires_receipt_link"]
        assert evidence() == before
        from app.db import SessionLocal
        from app.services.carton_supplier_portal import outstanding
        with SessionLocal() as db:
            assert outstanding(db, [line["id"] for line in order["lines"]]) == {}
        login_as(client, "admin")
        assert client.post(BASE + f"/internal/shipments/{shipment['id']}/receive", json=receive_payload(client, shipment)).status_code == 409
        options = client.get(BASE + f"/internal/shipments/{shipment['id']}/receipt-options", params={"factory_id": "huaxing"})
        assert options.status_code == 200 and options.json()[0]["id"] == receipt["id"], options.text
        body = link_body(shipment, receipt)
        linked = client.post(BASE + f"/internal/shipments/{shipment['id']}/link-receipt", json=body)
        assert linked.status_code == 200, linked.text
        assert linked.json()["linked_existing_receipt"] and linked.json()["receipt_id"] == receipt["id"]
        assert linked.json()["acceptance_date"] == "2026-09-21"
        assert evidence() == before
        replay_link = client.post(BASE + f"/internal/shipments/{shipment['id']}/link-receipt", json=body)
        assert replay_link.status_code == 200, replay_link.text
        assert client.post(BASE + f"/internal/shipments/{shipment['id']}/link-receipt", json=body | {"reason": "变更内容不能重复"}).status_code == 409
        supplier_login(client)
        replay = client.post(BASE + "/shipments/import-confirm", files={"file": ("dongkang.xlsx", content)}, data=data)
        assert replay.status_code == 200, replay.text
        changed_mode = dict(data, selections=json.dumps([{"factory_id": "huaxing", "delivery_note_no": "DK-IMPORT-01"}]))
        assert client.post(BASE + "/shipments/import-confirm", files={"file": ("dongkang.xlsx", content)}, data=changed_mode).status_code == 409
        public = client.get(BASE + "/workspace", params={"factory_id": "huaxing"}).json()["shipments"][0]
        assert public["linked_existing_receipt"] and "receipt_id" not in public
        assert all("unit_price" not in line and "location_allocations" not in line for line in public["acceptance_lines"])
        if rejected_last:
            assert Decimal(public["acceptance_lines"][-1]["rejected_quantity"]) == 10


@pytest.mark.parametrize("reverse_before_link", [False, True])
def test_reversal_reopens_original_note_then_corrects_once(monkeypatch, reverse_before_link):
    with make_client(monkeypatch) as client:
        setup_portal(client); order = accept_all(client)
        receipt = manual_receipt(client, order, note="DK-IMPORT-01")
        shipment, _, _ = late_note(client, order)
        login_as(client, "admin")
        if not reverse_before_link:
            linked = client.post(BASE + f"/internal/shipments/{shipment['id']}/link-receipt", json=link_body(shipment, receipt))
            assert linked.status_code == 200, linked.text
        response = client.post(f"/api/carton-procurement/receipts/{receipt['id']}/reverse", json={"factory_id": "huaxing", "expected_revision": receipt["revision"], "reason": "实际验收有误须更正"})
        assert response.status_code == 200, response.text
        if reverse_before_link:
            options = client.get(BASE + f"/internal/shipments/{shipment['id']}/receipt-options", params={"factory_id": "huaxing"}).json()
            assert options[0]["status"] == "REVERSED"
            body = link_body(shipment, receipt) | {"expected_receipt_revision": options[0]["revision"]}
            linked = client.post(BASE + f"/internal/shipments/{shipment['id']}/link-receipt", json=body)
            assert linked.status_code == 200, linked.text
        current = client.get(BASE + "/internal/workspace", params={"factory_id": "huaxing"}).json()["shipments"][0]
        assert current["requires_correction"] and not current["requires_receipt_link"]
        assert current["acceptance_history"] and current["acceptance_lines"] == []
        body = receive_payload(client, current) | {"correction_reason": "沿原凭证重新核实到货"}
        response = client.post(BASE + f"/internal/shipments/{shipment['id']}/receive", json=body)
        assert response.status_code == 200, response.text
        assert response.json()["receipt_id"] != receipt["id"]
        assert client.post(BASE + f"/internal/shipments/{shipment['id']}/receive", json=body).status_code == 200
        _, _, rows = evidence()
        assert all(Decimal(row["quantity"]) == 8 for row in rows)


def test_ordinary_new_delivery_requires_explicit_duplicate_check(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client); order = accept_all(client)
        manual_receipt(client, order)
        supplier_login(client)
        shipment = client.post(BASE + "/shipments", json=ship_payload(order)).json()
        login_as(client, "admin")
        body = receive_payload(client, shipment)
        url = BASE + f"/internal/shipments/{shipment['id']}/receive"
        blocked = client.post(url, json=body)
        assert blocked.status_code == 409 and "已入库记录" in blocked.text, blocked.text
        received = client.post(url, json=body | {"new_delivery_confirmation": True})
        assert received.status_code == 200, received.text


def test_link_scope_versions_and_one_receipt_one_note(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client); order = accept_all(client)
        receipt = manual_receipt(client, order)
        shipment, _, _ = late_note(client, order)
        url = BASE + f"/internal/shipments/{shipment['id']}/link-receipt"
        body = link_body(shipment, receipt)
        assert client.post(url, json=body).status_code == 403  # Supplier cannot choose a warehouse receipt.
        login_as(client, "admin")
        assert client.post(url, json=body | {"factory_id": "huakang-a"}).status_code in (403, 404)
        assert client.post(url, json=body | {"expected_revision": 999}).status_code == 409
        assert client.post(url, json=body | {"expected_receipt_revision": 999}).status_code == 409
        supplier_login(client)
        second = client.post(BASE + "/shipments", json=ship_payload(order) | {"delivery_note_no": "NEW-NOTE"}).json()
        login_as(client, "admin")
        assert client.post(url, json=body).status_code == 200
        other = client.post(BASE + f"/internal/shipments/{second['id']}/link-receipt", json=link_body(second, receipt))
        assert other.status_code == 409, other.text
        duplicates = client.get(BASE + f"/internal/shipments/{second['id']}/receipt-options", params={"factory_id": "huaxing"}).json()
        assert duplicates[0]["already_linked"]
        blocked = client.post(BASE + f"/internal/shipments/{second['id']}/receive", json=receive_payload(client, second))
        assert blocked.status_code == 409 and "已入库记录" in blocked.text, blocked.text


def test_completed_order_can_attach_original_file(monkeypatch):
    from io import BytesIO
    from openpyxl import load_workbook
    with make_client(monkeypatch) as client:
        setup_portal(client); order = accept_all(client)
        receipt = manual_receipt(client, order, full=True)
        workbook = load_workbook(BytesIO(_dongkang_delivery_file(order)))
        for index, line in enumerate(order["lines"], 2):
            workbook.active.cell(index, 8).value = float(line["required_quantity"])
        output = BytesIO(); workbook.save(output)
        shipment, _, _ = late_note(client, order, content=output.getvalue())
        login_as(client, "admin")
        response = client.post(BASE + f"/internal/shipments/{shipment['id']}/link-receipt", json=link_body(shipment, receipt))
        assert response.status_code == 200, response.text
        work = client.get(BASE + "/internal/workspace", params={"factory_id": "huaxing"}).json()
        assert work["orders"][0]["status"] == "COMPLETED"
        assert all(Decimal(line["in_transit_quantity"]) == 0 for line in work["orders"][0]["lines"])


@pytest.mark.parametrize("legacy_date", [False, True])
def test_locked_month_rejects_link_then_reopened_month_retains_original_period(monkeypatch, legacy_date):
    from io import BytesIO
    from openpyxl import load_workbook
    with make_client(monkeypatch) as client:
        setup_portal(client); order = accept_all(client)
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonReceipt
        from app.models.carton_supplier_settlement import CartonSupplierSettlement
        from app.services.carton_settlement_collaboration import build
        from app.services.carton_supplier_settlement import sources
        receipt = manual_receipt(client, order)
        workbook = load_workbook(BytesIO(_dongkang_delivery_file(order)))
        for index in range(2, len(order["lines"])+2): workbook.active.cell(index, 2).value = "2026-10-05"
        output = BytesIO(); workbook.save(output)
        shipment, _, _ = late_note(client, order, content=output.getvalue())
        if legacy_date:
            with SessionLocal() as db:
                original = db.get(CartonReceipt, receipt["id"])
                original.acceptance_date = None
                original.confirmed_at = "2026-09-20T17:30:00Z"  # September 21 in Shanghai.
                db.commit()
        before = evidence()
        with SessionLocal() as db:
            supplier = db.get(CartonReceipt, receipt["id"]).supplier_id
            db.add(CartonSupplierSettlement(id="CLOSED-MONTH", factory_id="huaxing", supplier_id=supplier,
                period="2026-09", currency="CNY", version=1, revision=1, status="CONFIRMED", source_fingerprint="0"*64,
                created_at="2026-10-01", updated_at="2026-10-01", created_by="admin"))
            db.commit()
        login_as(client, "admin")
        url = BASE + f"/internal/shipments/{shipment['id']}/link-receipt"
        body = link_body(shipment, receipt)
        blocked = client.post(url, json=body)
        assert blocked.status_code == 409 and "月份已确认" in blocked.text, blocked.text
        assert evidence() == before
        with SessionLocal() as db:
            db.get(CartonSupplierSettlement, "CLOSED-MONTH").status = "DRAFT"; db.commit()
        linked = client.post(url, json=body)
        assert linked.status_code == 200 and linked.json()["acceptance_date"] == "2026-09-21", linked.text
        with SessionLocal() as db:
            rows, base, _, _ = sources(db, "huaxing", supplier, "2026-09", "CNY")
            collaboration = build(db, "huaxing", supplier, "2026-09", "CNY", rows, base)
            assert len(collaboration["sources"]) == len(order["lines"])
            assert all(row["supplier_delivery"]["document_no"] == "DK-IMPORT-01" for row in collaboration["sources"])
            assert not sources(db, "huaxing", supplier, "2026-10", "CNY")[0]
            if legacy_date:
                assert db.get(CartonReceipt, receipt["id"]).acceptance_date is None


def test_voided_drafts_and_unrelated_reversals_are_never_recovery_candidates(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client); order = accept_all(client)
        login_as(client, "admin")
        draft = client.post("/api/carton-procurement/receipts", json={"factory_id": "huaxing", "delivery_note_no": "VOIDED-DRAFT",
            "delivery_date": "2026-09-21", "lines": [{"order_line_id": line["id"], "delivered_quantity": 10, "received_quantity": 10, "unit_price": 2} for line in order["lines"]]})
        assert draft.status_code == 201, draft.text
        def reverse(receipt):
            response = client.post(f"/api/carton-procurement/receipts/{receipt['id']}/reverse", json={"factory_id": "huaxing", "expected_revision": receipt["revision"], "reason": "复核原始记录需要冲销"})
            assert response.status_code == 200, response.text
        reverse(draft.json())
        unrelated = manual_receipt(client, order, note="UNRELATED-OLD")
        reverse(unrelated)
        original = manual_receipt(client, order, note="MANUAL-CURRENT")
        shipment, _, _ = late_note(client, order)
        login_as(client, "admin")
        reverse(original)
        url = BASE + f"/internal/shipments/{shipment['id']}"
        candidates = client.get(url + "/receipt-options", params={"factory_id": "huaxing"}).json()
        assert [row["id"] for row in candidates] == [original["id"]]
        for bad_receipt in (draft.json(), unrelated):
            bad = client.post(url + "/link-receipt", json=link_body(shipment, bad_receipt))
            assert bad.status_code == 409, bad.text
