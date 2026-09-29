"""Synthetic closure; no device or production-store acceptance claims."""
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from random import Random
from uuid import uuid4
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from test_uv_ops import client, post, read, payload, setup_task, confirm, user_with
from app.models import uv_operations as m
from app.services.uv_operations import common as c, exports, payroll, live


def block(data, *, batch=None, machine=None, **extra):
    batch, machine = batch or data['batch'], machine or data['machine']
    return dict(task_id=data['task']['id'], batch_id=batch['id'], batch_version=batch['version'], machine_id=machine['id'], task_version=data['task']['version'], machine_version=machine['version'], block_version=0, start_at='2026-09-23T08:00:00+08:00', end_at='2026-09-23T09:00:00+08:00', **extra)


def request(client, path, **body):
    return client.post('/api/uv-operations/'+path,json=payload(**body))


def participant(client, data, employee='worker'):
    return post(client,'participations',shift_id=data['shift']['id'],task_id=data['task']['id'],employee_id=employee,employee_name='合成员工',start_at='2026-09-23T08:00:00+08:00',end_at='2026-09-23T09:00:00+08:00')


def test_manual_estimate_and_independent_parallel_batches(client):
    data=setup_task(client)
    # Represent the documented unknown-cycle frozen task explicitly.
    with Session(client.uv_engine) as db:
        task=c.get(db,m.UvOpsTask,data['task']['id'])
        task.process_snapshot=dict(task.process_snapshot,cycle_seconds=None)
        db.commit()
    split=post(client,'batches/'+data['batch']['id']+'/split',expected_version=1,quantity=60,reason='合成双机实体拆批')
    data['task']=read(client,'tasks/'+data['task']['id'])
    one,two=split['batches']
    machine=post(client,'machines',code='PARALLEL',name='合成并行机台',width_mm='1000',height_mm='800',ink_family='synthetic-hard',capability_evidence='合成能力验收')
    fixture=post(client,'fixtures',code='SECOND-PHYSICAL',revision=1,slots=12,width_mm='100',height_mm='80')
    first=block(data,batch=one)
    assert not post(client,'schedule/preview',blocks=[first])['can_commit']
    first.update(manual_estimated_seconds='300',estimate_reason='合成首件计时估计')
    second=block(data,batch=two,machine=machine,manual_estimated_seconds='300',estimate_reason='合成首件计时估计')
    assert not post(client,'schedule/preview',blocks=[first,second])['can_commit']
    second.update(fixture_id=fixture['id'],fixture_version=fixture['version'],fixture_evidence='独立实体治具首件验证')
    result=post(client,'schedule/commit',blocks=[first,second])
    assert len(result['blocks'])==2 and all(x['estimate_basis']=='manual' for x in result['blocks'])
    assert read(client,'tasks/'+data['task']['id'])['process_snapshot']['cycle_seconds'] is None
    execution=post(client,'schedule/'+result['blocks'][0]['id']+'/start',expected_version=1,shift_id=data['shift']['id'],started_at='2026-09-23T08:00:00+08:00',evidence='现场人工开工合成证据')
    data['task']=read(client,'tasks/'+data['task']['id'])
    second.update(task_version=data['task']['version'],block_version=1,start_at='2026-09-23T10:00:00+08:00',end_at='2026-09-23T11:00:00+08:00')
    assert post(client,'schedule/preview',blocks=[second])['can_commit']
    first.update(task_version=data['task']['version'],block_version=1)
    assert not post(client,'schedule/preview',blocks=[first])['can_commit']
    assert request(client,'batches/'+one['id']+'/split',expected_version=one['version'],quantity=10,reason='已开工不得再拆批').status_code==409
    gaps=read(client,'shifts/'+data['shift']['id']+'/gaps')
    assert execution['id'] in gaps['running_executions'] and len(gaps['unreported_plans'])==2


def test_fixture_revisions_conflict_and_batch_duration(client):
    data=setup_task(client)
    split=post(client,'batches/'+data['batch']['id']+'/split',expected_version=1,quantity=12,reason='按实体批次核算节拍')
    data['task']=read(client,'tasks/'+data['task']['id'])
    first=block(data,batch=split['batches'][0]);first['end_at']='2026-09-23T08:00:30+08:00'
    assert Decimal(post(client,'schedule/preview',blocks=[first])['blocks'][0]['estimated_seconds'])==30
    fixture=post(client,'fixtures',code=data['fixture']['code'],revision=2,slots=12,width_mm='100',height_mm='80')
    machine=post(client,'machines',code='REVISION',name='另机',width_mm='1000',height_mm='800',ink_family='synthetic-hard',capability_evidence='合成核对能力')
    second=block(data,batch=split['batches'][1],machine=machine,fixture_id=fixture['id'],fixture_version=1,fixture_evidence='同实体的另一修订')
    assert not post(client,'schedule/preview',blocks=[first,second])['can_commit']
    first['manual_estimated_seconds']='10'
    assert not post(client,'schedule/preview',blocks=[first])['can_commit']


def test_same_batch_concurrent_plan_and_fixed_reshape(client):
    data=setup_task(client)
    item=block(data,fixed=True)
    with ThreadPoolExecutor(2) as pool:
        responses=list(pool.map(lambda _:request(client,'schedule/commit',blocks=[item]),range(2)))
    assert sorted(x.status_code for x in responses)==[200,409]
    plan=read(client,'schedule')[0]
    assert request(client,'batches/'+data['batch']['id']+'/split',expected_version=1,quantity=50,reason='合成拆批应被固定阻止').status_code==409
    post(client,'schedule/'+plan['id']+'/unlock',expected_version=plan['version'],reason='合成重排先解锁计划')
    post(client,'batches/'+data['batch']['id']+'/split',expected_version=1,quantity=50,reason='合成拆批后旧计划撤销')
    assert read(client,'schedule')[0]['status']=='cancelled'


def test_cross_shift_manual_execution_finishes_without_reopening(client):
    data=setup_task(client)
    plan=post(client,'schedule/commit',blocks=[block(data)])['blocks'][0]
    run=post(client,'schedule/'+plan['id']+'/start',expected_version=1,shift_id=data['shift']['id'],started_at='2026-09-23T15:50:00+08:00',evidence='合成跨班持续生产')
    assert request(client,'shifts/'+data['shift']['id']+'/close',expected_version=1).status_code==409
    post(client,'shifts/'+data['shift']['id']+'/close',expected_version=1,wip_handover='合成在制批次交接晚班继续')
    next_shift=post(client,'shifts',name='合成晚班',business_date='2026-09-23',start_at='2026-09-23T16:00:00+08:00',end_at='2026-09-24T00:00:00+08:00')
    finished=post(client,'executions/'+run['id']+'/finish',expected_version=1,shift_id=next_shift['id'],ended_at='2026-09-23T17:00:00+08:00',evidence='合成晚班人工完工核对')
    assert finished['ended_shift_id']==next_shift['id'] and finished['active_key'] is None
    assert next(x for x in read(client,'shifts') if x['id']==data['shift']['id'])['status']=='closed'
    assert read(client,'production')==[] and read(client,'runs')==[]


def test_correction_quality_payroll_and_participation_chain(client):
    data=setup_task(client,quantity=20)
    result=confirm(client,data,good=10,rework=0,pending=10)
    person=participant(client,data)
    quality=post(client,'quality/resolve',batch_id=data['batch']['id'],shift_id=data['shift']['id'],expected_version=result['batch']['version'],quantity=10,disposition='good',evidence='合成待检转良品')['entry']
    wage=post(client,'wages/confirm',task_id=data['task']['id'],shift_id=data['shift']['id'])['accrual']
    assert request(client,'quality/'+quality['id']+'/reverse',expected_version=1,reason='合成错误品质核对').status_code==409
    operation=payload(expected_version=wage['version'],reason='合成核准后发现品质误录')
    url='/api/uv-operations/wages/'+wage['id']+'/reverse'
    assert client.post(url,json=operation).status_code==client.post(url,json=operation).status_code==200
    post(client,'quality/'+quality['id']+'/reverse',expected_version=1,reason='合成品质结论录入错误')
    post(client,'participations/'+person['id']+'/reverse',expected_version=1,reason='合成员工时段录入错误')
    assert read(client,'wages')==read(client,'wage-allocations')==read(client,'quality')==read(client,'participations')==[]
    batch=read(client,'tasks/'+data['task']['id'])['batches'][0]
    assert batch['pending']==10 and batch['good']==10
    post(client,'production/'+result['entry']['id']+'/reverse',expected_version=batch['version'],reason='合成重新核数纠正报产')
    data['batch']=read(client,'tasks/'+data['task']['id'])['batches'][0]
    confirm(client,data,good=19,rework=0,scrap=1)
    participant(client,data)
    new=post(client,'wages/confirm',task_id=data['task']['id'],shift_id=data['shift']['id'])
    assert Decimal(new['accrual']['payroll_amount'])==Decimal('0.29')
    assert new['accrual']['id']!=wage['id'] and sum(Decimal(x['payroll_amount']) for x in new['allocations'])==Decimal('0.29')
    history=client.get('/api/uv-operations/wages',params=dict(factory_id=m.FACTORY,history=True)).json()['data']
    assert len(history)==2 and sum(x['active_key']==1 for x in history)==1


def test_rework_correction_can_terminate_fully_reversed_child(client):
    data=setup_task(client,quantity=12,passes=3)
    initial=confirm(client,data,good=0,rework=0,pending=12)
    quality=post(client,'quality/resolve',batch_id=data['batch']['id'],shift_id=data['shift']['id'],expected_version=initial['batch']['version'],quantity=12,disposition='rework',evidence='合成误判返工')
    rework=post(client,'rework-tasks',batch_id=data['batch']['id'],expected_version=quality['batch']['version'],code='CORRECT-REWORK',quantity=12)
    child=dict(data,task=rework,batch=quality['batch'])
    item=block(child);item['end_at']='2026-09-23T08:00:30+08:00'
    assert post(client,'schedule/preview',blocks=[item])['can_commit']
    result=confirm(client,child,good=12,rework=0,source_bucket='rework')
    post(client,'production/'+result['entry']['id']+'/reverse',expected_version=result['batch']['version'],reason='合成冲销返工加工核数')
    child['task']=read(client,'tasks/'+rework['id'])
    post(client,'rework-tasks/'+rework['id']+'/cancel',expected_version=child['task']['version'],reason='合成所有报产已冲销终止')
    post(client,'quality/'+quality['entry']['id']+'/reverse',expected_version=1,reason='合成撤销最初错误品质判断')
    assert read(client,'tasks/'+data['task']['id'])['batches'][0]['pending']==12


def test_expense_reverse_preserves_closed_month_snapshot(client):
    data=setup_task(client,quantity=12)
    confirm(client,data,good=12,rework=0)
    participant(client,data)
    post(client,'wages/confirm',task_id=data['task']['id'],shift_id=data['shift']['id'])
    expense=post(client,'expenses',task_id=data['task']['id'],category='direct',cost_amount='1.23',currency='CNY',business_date='2026-09-23',evidence='合成费用凭证')
    post(client,'shifts/'+data['shift']['id']+'/close',expected_version=1)
    period=read(client,'periods')[0]
    closed=post(client,'periods/2026-09/close',expected_version=period['version'],reason='合成月度结算验收')
    assert request(client,'expenses/'+expense['id']+'/reverse',expected_version=1,reason='合成封账应当阻止纠错').status_code==409
    post(client,'periods/2026-09/reopen',expected_version=closed['period']['version'],reason='合成经复核重开账期纠错')
    post(client,'expenses/'+expense['id']+'/reverse',expected_version=1,reason='合成错误成本凭证冲销')
    assert read(client,'expenses')==[]
    assert read(client,'report-snapshots')[0]['report']==closed['snapshot']['report']


def workbook(client, kind, ids):
    job=post(client,'exports',kind=kind,selected_ids=ids,start_date='2026-09-01',end_date='2026-09-30')
    with Session(client.uv_engine) as db:
        row=c.get(db,m.UvOpsExportJob,job['id']);exports.build(db,client.uv_auth['user'],row)
        return load_workbook(BytesIO(row.artifact))


def test_exports_linked_details_and_no_false_selected_reports(client):
    data=setup_task(client,quantity=12)
    result=confirm(client,data,good=10,rework=0,pending=2)
    quality=post(client,'quality/resolve',batch_id=data['batch']['id'],shift_id=data['shift']['id'],expected_version=result['batch']['version'],quantity=2,disposition='good',evidence='=FORMULA-SAFE 品质证据')
    participant(client,data,employee='=SYNTHETIC-FORMULA')
    wage=post(client,'wages/confirm',task_id=data['task']['id'],shift_id=data['shift']['id'])
    handover=post(client,'handovers',batch_id=data['batch']['id'],expected_version=quality['batch']['version'],quantity=12,target='合成下工序',business_date='2026-09-23')
    post(client,'handovers/'+handover['id']+'/receive',expected_version=1,quantity=12,business_date='2026-09-23',evidence='=FORMULA-SAFE 签收证据')
    for kind,ids,sheet in [('production',[result['entry']['id']],'品质处置明细'),('wages',[wage['accrual']['id']],'员工工资分配'),('handovers',[handover['id']],'签收拒收退回')]:
        book=workbook(client,kind,ids)
        assert book[sheet].max_row==2
        assert all(cell.data_type!='f' for row in book[sheet] for cell in row)
    for kind in ('daily','monthly'):
        assert request(client,'exports',kind=kind,selected_ids=[data['task']['id']],start_date='2026-09-01',end_date='2026-09-30').status_code==422


def test_workspace_cache_revisions_projection_and_pool_concurrency(client):
    data=setup_task(client)
    def snapshot(_):
        with Session(client.uv_engine) as db:
            return live.response(db,client.uv_auth['user'])
    with ThreadPoolExecutor(10) as pool: values=list(pool.map(snapshot,range(20)))
    assert all(x['data']['collection_totals']['tasks']==1 for x in values)
    # Per-user projection cannot contaminate the shared raw snapshot.
    with Session(client.uv_engine) as db:
        limited=live.response(db,user_with('read'))
        assert 'cost_price_snapshot' not in limited['data']['tasks'][0]
    post(client,'machines',code='NEW-REV',name='合成缓存失效')
    current=snapshot(0)
    assert current['meta']['view_revision']>values[0]['meta']['view_revision'] and len(current['data']['machines'])==2


def test_randomized_money_conservation_and_determinism():
    rng=Random(20260929)
    for _ in range(1000):
        cents=rng.randrange(0,10_000_000)
        weights={f'worker-{i}':Decimal(rng.randrange(1,86401))/Decimal(rng.randrange(1,10)) for i in range(rng.randrange(1,21))}
        amount=Decimal(cents)/100
        split=payroll.split_money(amount,weights)
        assert sum(split.values())==amount and all(x>=0 and x.as_tuple().exponent>=-2 for x in split.values())
        assert split==payroll.split_money(amount,dict(reversed(list(weights.items()))))


def test_schedule_window_paginates_complete_thirty_days(client):
    data=setup_task(client)
    with Session(client.uv_engine) as db:
        task=c.get(db,m.UvOpsTask,data['task']['id'])
        values={key:getattr(task,key) for key in ('demand_id','process_version_id','file_version_id','quantity','product_snapshot','process_snapshot','cost_price_snapshot','payroll_policy_snapshot')}
        for index in range(605):
            row=c.add(db,m.UvOpsTask,code=f'WINDOW-{index}',**values)
            batch=c.add(db,m.UvOpsBatch,task_id=row.id,quantity=120,remaining=120)
            start=datetime.fromisoformat('2026-09-01T00:00:00+08:00')+timedelta(days=index%30,minutes=(index//30)*10)
            c.add(db,m.UvOpsScheduleBlock,task_id=row.id,batch_id=batch.id,machine_id=data['machine']['id'],fixture_id=data['fixture']['id'],start_at=c.ts(start),end_at=c.ts(start+timedelta(minutes=5)))
        db.commit()
    params=dict(factory_id=m.FACTORY,start_at='2026-09-01T00:00:00+08:00',end_at='2026-10-01T00:00:00+08:00',limit=500)
    first=client.get('/api/uv-operations/schedule/window',params=params).json()
    second=client.get('/api/uv-operations/schedule/window',params=params|dict(cursor=first['pagination']['next_cursor'])).json()
    assert first['pagination']['total']==605 and len(first['data'])==500 and len(second['data'])==105
    assert len({x['id'] for x in first['data']+second['data']})==605 and not second['pagination']['has_more']
    assert client.get('/api/uv-operations/schedule/window',params=params|dict(end_at='2026-11-01T00:00:00+08:00')).status_code==422


def test_run_cost_pool_reversal_is_atomic_and_reconfirmable(client):
    from test_uv_ops_edges import setup_agent,event
    data=setup_task(client,quantity=12)
    machine,agent,source,headers=setup_agent(client)
    response=client.post('/api/internal/uv-agent/events/batch',headers=headers,json=dict(schema_version=1,batch_id=uuid4().hex,events=[event(machine,source,kind='job_complete')]))
    assert response.json()['results'][0]['status']=='persisted'
    run=read(client,'runs')[0]
    matched=post(client,'runs/'+run['id']+'/match',expected_version=run['version'],evidence='合成确认运行分配',allocations=[dict(task_id=data['task']['id'],batch_id=data['batch']['id'],pass_index=1,slots=list(range(1,13)),full_boards=1,tail_pieces=0,share='1')])['run']
    pool=post(client,'run-costs',run_id=run['id'],expected_version=matched['version'],cost_amount='30.01',currency='CNY',business_date='2026-09-23',evidence='合成运行成本池')
    expense=read(client,'expenses')[0]
    assert request(client,'expenses/'+expense['id']+'/reverse',expected_version=1,reason='不得单独撤销成本子项').status_code==409
    post(client,'run-costs/'+pool['id']+'/reverse',expected_version=1,reason='合成运行成本池录入错误')
    assert read(client,'expenses')==read(client,'run-costs')==[]
    replacement=post(client,'run-costs',run_id=run['id'],expected_version=matched['version'],cost_amount='10.02',currency='CNY',business_date='2026-09-23',evidence='合成重新核准成本池')
    assert replacement['id']!=pool['id'] and sum(Decimal(x['cost_amount']) for x in read(client,'expenses'))==Decimal('10.02')


def test_export_cross_day_quality_and_receipt_use_fact_dates(client):
    data=setup_task(client,quantity=12)
    result=confirm(client,data,good=10,rework=0,pending=2)
    shift=post(client,'shifts',name='合成次日班',business_date='2026-09-24',start_at='2026-09-24T08:00:00+08:00',end_at='2026-09-24T16:00:00+08:00')
    quality=post(client,'quality/resolve',batch_id=data['batch']['id'],shift_id=shift['id'],quantity=2,disposition='good',expected_version=result['batch']['version'],evidence='合成次日品质结论')
    handover=post(client,'handovers',batch_id=data['batch']['id'],quantity=10,expected_version=quality['batch']['version'],target='合成跨日签收',business_date='2026-09-23')
    post(client,'handovers/'+handover['id']+'/receive',expected_version=1,quantity=10,business_date='2026-09-24',evidence='合成次日点收')
    for kind,title,source in [('production','品质处置明细','品质来源核数'),('handovers','签收拒收退回','关联交接原单')]:
        job=post(client,'exports',kind=kind,selected_ids=[],start_date='2026-09-24',end_date='2026-09-24')
        with Session(client.uv_engine) as db:
            row=c.get(db,m.UvOpsExportJob,job['id']);exports.build(db,client.uv_auth['user'],row)
            book=load_workbook(BytesIO(row.artifact))
            assert book[title].max_row==2 and book[source].max_row==2
            headers=[x.value for x in book[title][1]]
            assert book[title].cell(2,headers.index('business_date')+1).value=='2026-09-24'


def test_randomized_output_corrections_and_inventory_replay(client):
    rng=Random(929)
    data=setup_task(client,quantity=200)
    total_good=total_scrap=0
    for index in range(15):
        qty=rng.randrange(1,10);good=rng.randrange(qty+1)
        result=confirm(client,data,good=good,rework=0,scrap=qty-good)
        if index%3==0:
            post(client,'production/'+result['entry']['id']+'/reverse',expected_version=result['batch']['version'],reason='合成随机序列纠正核数')
            data['batch']=read(client,'tasks/'+data['task']['id'])['batches'][0]
            good=rng.randrange(qty+1);result=confirm(client,data,good=good,rework=0,scrap=qty-good)
        total_good+=good;total_scrap+=qty-good;data['batch']=result['batch']
        assert data['batch']['good']==total_good and data['batch']['scrap']==total_scrap
        assert data['batch']['remaining']+total_good+total_scrap==200
    sku=post(client,'ink/skus',code='RANDOM-INK',supplier='合成',model='合成',ink_family='synthetic',color='C',capacity_ml='1000')
    stock=post(client,'ink/movements',sku_id=sku['id'],location='STORE',kind='receipt',quantity_ml='1000',cost_value='100',currency='CNY',business_date='2026-09-23',evidence='合成随机库存起点')['balance']
    remaining=Decimal(1000)
    for index in range(15):
        qty=Decimal(rng.randrange(1,30))
        body=payload(sku_id=sku['id'],location='STORE',kind='consume',quantity_ml=str(qty),currency='CNY',business_date='2026-09-23',expected_version=stock['version'],evidence='合成库存幂等核对')
        url='/api/uv-operations/ink/movements'
        first=client.post(url,json=body);replay=client.post(url,json=body)
        assert first.status_code==replay.status_code==200 and first.json()['data']==replay.json()['data']
        value=first.json()['data'];stock=value['balance'];remaining-=qty
        if index%4==0:
            stock=post(client,'ink/movements',sku_id=sku['id'],location='STORE',kind='reverse',quantity_ml=str(qty),currency='CNY',business_date='2026-09-23',expected_version=stock['version'],reversal_of=value['entry']['id'],evidence='合成原成本退回核对')['balance'];remaining+=qty
        assert Decimal(stock['quantity_ml'])==remaining and Decimal(stock['cost_value'])==remaining/10


def test_postgres_plan_match_ingest_lock_order_no_cycle(client):
    import pytest
    from threading import Event,Thread,current_thread
    from sqlalchemy import event as sql_event
    from test_uv_ops_edges import setup_agent,event
    from app.services.uv_operations import planning,ingest
    from app.schemas import uv_operations as s
    if client.uv_engine.dialect.name!='postgresql':
        pytest.skip('PostgreSQL row-lock interleaving; SQLite covered by command serialization tests')
    data=setup_task(client)
    machine,agent,source,headers=setup_agent(client)
    with Session(client.uv_engine) as db:
        row=c.get(db,m.UvOpsMachine,machine['id']);row.width_mm=1000;row.height_mm=800;row.ink_family='synthetic-hard';row.capability_evidence='合成机台能力'
        db.commit();data['machine']=c.record(row)
    raw=event(machine,source)
    with Session(client.uv_engine) as db: ingest.event(db,agent['agent']['id'],raw)
    run=read(client,'runs')[0]
    plan_body=s.SchedulePreview(**payload(blocks=[block(data)]))
    match_body=s.RunMatch(**payload(expected_version=run['version'],evidence='合成并发确认槽位',allocations=[dict(task_id=data['task']['id'],batch_id=data['batch']['id'],pass_index=1,slots=list(range(1,13)),full_boards=10,tail_pieces=0,share='1')]))
    task_locked,match_attempt,machine_locked,release_plan,release_event=[Event() for _ in range(5)]
    results={}
    def before(conn,cursor,statement,params,context,many):
        if current_thread().name=='uv-match' and 'FROM uv_ops_tasks' in statement and 'FOR UPDATE' in statement:
            match_attempt.set()
    def after(conn,cursor,statement,params,context,many):
        if current_thread().name=='uv-plan' and 'FROM uv_ops_tasks' in statement and 'FOR UPDATE' in statement and not task_locked.is_set():
            task_locked.set();assert release_plan.wait(10)
        if current_thread().name=='uv-ingest' and 'FROM uv_ops_machines' in statement and 'FOR UPDATE' in statement:
            machine_locked.set();assert release_event.wait(10)
    def work(name,handler):
        try:
            with Session(client.uv_engine) as db: results[name]=handler(db)
        except Exception as error: results[name]=error
    user=client.uv_auth['user']
    sql_event.listen(client.uv_engine,'before_cursor_execute',before);sql_event.listen(client.uv_engine,'after_cursor_execute',after)
    threads=[]
    try:
        for name,handler,ready in [
            ('uv-plan',lambda db:c.command(db,user,'test-plan',plan_body,('plan_write',),lambda:planning.commit(db,user,plan_body)),task_locked),
            ('uv-match',lambda db:c.command(db,user,'test-match',match_body,('production_write',),lambda:ingest.match_run(db,user,run['id'],match_body)),match_attempt),
            ('uv-ingest',lambda db:ingest.event(db,agent['agent']['id'],event(machine,source,2,count='1')),machine_locked)]:
            thread=Thread(name=name,target=work,args=(name,handler));thread.start();threads.append(thread)
            assert ready.wait(10),name
        release_plan.set();release_event.set()
        for thread in threads: thread.join(10);assert not thread.is_alive()
        assert not isinstance(results['uv-plan'],Exception) and not isinstance(results['uv-ingest'],Exception),results
        if isinstance(results['uv-match'],c.DomainError):
            assert results['uv-match'].body['code']=='version_conflict'
            fresh=read(client,'runs')[0]
            post(client,'runs/'+run['id']+'/match',expected_version=fresh['version'],evidence=match_body.evidence,allocations=[x.model_dump(mode='json') for x in match_body.allocations])
        else:
            assert not isinstance(results['uv-match'],Exception),results
        assert len(read(client,'schedule'))==1 and len(read(client,'run-allocations'))==1
    finally:
        release_plan.set();release_event.set()
        for thread in threads: thread.join(15)
        sql_event.remove(client.uv_engine,'before_cursor_execute',before);sql_event.remove(client.uv_engine,'after_cursor_execute',after)


def test_running_execution_allows_future_plan_but_not_actual_start(client,monkeypatch):
    monkeypatch.setattr(m,'now',lambda:'2026-09-23T00:30:00+00:00')
    data=setup_task(client)
    plan=post(client,'schedule/commit',blocks=[block(data)])['blocks'][0]
    post(client,'schedule/'+plan['id']+'/start',expected_version=1,shift_id=data['shift']['id'],started_at='2026-09-23T08:00:00+08:00',evidence='合成实际开工证据')
    next_data=setup_task(client)
    next_data['machine']=data['machine']
    future=block(next_data)
    assert not post(client,'schedule/preview',blocks=[future])['can_commit']
    monkeypatch.setattr(m,'now',lambda:'2026-09-23T01:23:45.123456+00:00')
    suggested=client.post('/api/uv-operations/schedule/recommendations',json=dict(factory_id='huakang-a',task_id=next_data['task']['id'],batch_id=next_data['batch']['id'],earliest_at='2026-09-23T09:00:00+08:00')).json()['data']
    chosen=next(row for row in suggested if row['machine_id']==data['machine']['id'])
    assert chosen['start_at']=='2026-09-23T01:24:00.000000+00:00'
    adopted=dict(future,start_at=chosen['start_at'],end_at=chosen['end_at'])
    assert post(client,'schedule/preview',blocks=[adopted])['can_commit']
    future.update(start_at='2026-09-24T08:00:00+08:00',end_at='2026-09-24T09:00:00+08:00')
    preview=post(client,'schedule/preview',blocks=[future]);assert preview['can_commit'] and preview['warnings']
    suggestions=client.post('/api/uv-operations/schedule/recommendations',json=dict(factory_id='huakang-a',task_id=next_data['task']['id'],batch_id=next_data['batch']['id'],earliest_at='2026-09-24T08:00:00+08:00')).json()['data']
    assert any(row['machine_id']==data['machine']['id'] and row['compatible'] for row in suggestions)
    saved=post(client,'schedule/commit',blocks=[future])['blocks'][0]
    shift=post(client,'shifts',name='未来合成班',business_date='2026-09-24',start_at='2026-09-24T08:00:00+08:00',end_at='2026-09-24T16:00:00+08:00')
    assert request(client,'schedule/'+saved['id']+'/start',expected_version=1,shift_id=shift['id'],started_at='2026-09-24T08:00:00+08:00',evidence='不能绕过实际机台占用').status_code==409


def test_ink_date_order_source_target_and_corrected_task_cancellation(client):
    data=setup_task(client)
    sku=post(client,'ink/skus',code='DATED',supplier='合成',model='合成',ink_family='synthetic',color='C',capacity_ml='1000')
    def move(**values):return dict(sku_id=sku['id'],location='A',currency='CNY',quantity_ml='10',evidence='合成跨日库存核对',**values)
    stock=post(client,'ink/movements',**move(kind='receipt',cost_value='10',business_date='2026-09-29'))['balance']
    result=request(client,'ink/movements',**move(kind='consume',business_date='2026-09-28',expected_version=stock['version']))
    assert result.status_code==409 and result.json()['code']=='inventory_date_order'
    used=post(client,'ink/movements',**move(kind='consume',task_id=data['task']['id'],business_date='2026-09-29',expected_version=stock['version']))
    assert request(client,'tasks/'+data['task']['id']+'/cancel',expected_version=1,reason='合成误建任务取消').status_code==409
    assert request(client,'ink/movements',**move(kind='reverse',reversal_of=used['entry']['id'],business_date='2026-08-31',expected_version=used['balance']['version'])).status_code==409
    restored=post(client,'ink/movements',**move(kind='reverse',reversal_of=used['entry']['id'],business_date='2026-09-29',expected_version=used['balance']['version']))
    assert post(client,'tasks/'+data['task']['id']+'/cancel',expected_version=1,reason='合成误耗已完整冲销')['demand']['allocated']==0
    target=post(client,'ink/movements',**(move(kind='receipt',cost_value='10',business_date='2026-10-01')|{'location':'B'}))
    assert request(client,'ink/movements',**move(kind='transfer',target_location='B',business_date='2026-09-30',expected_version=restored['balance']['version'])).status_code==409
    moved=post(client,'ink/movements',**move(kind='transfer',target_location='B',business_date='2026-10-01',expected_version=restored['balance']['version']))
    assert Decimal(moved['target']['quantity_ml'])==20


def test_live_current_run_follows_source_identity_rebinding_and_history_window(client):
    from test_uv_ops_edges import setup_agent,event
    from app.services.uv_operations import ingest
    machine,agent,source,headers=setup_agent(client)
    def upload(raw):
        with Session(client.uv_engine) as db:return ingest.event(db,agent['agent']['id'],raw)
    upload(event(machine,source))
    with Session(client.uv_engine) as db:
        c.get(db,m.UvOpsAgent,agent['agent']['id']).last_seen_at=m.now()
        db.scalar(c.query(m.UvOpsSettings)).stale_seconds=10000000
        original=db.scalar(c.query(m.UvOpsRun));original_id=original.id
        for i in range(105):c.add(db,m.UvOpsRun,machine_id=machine['id'],binding_id=source['id'],native_job_id='historical'+str(i),native_identity='history'+str(i),started_at='2026-09-01T00:00:00+00:00',ended_at='2026-09-01T00:01:00+00:00',last_observed_at='2026-09-01T00:01:00+00:00')
        db.commit()
    snapshot=read(client,'workspace');assert snapshot['machines'][0]['current_run_id']==original_id and any(x['id']==original_id for x in snapshot['runs'])
    post(client,'sources/'+source['id']+'/unbind',expected_version=source['version'],reason='合成来源换流确认')
    assert read(client,'workspace')['machines'][0]['current_run_id'] is None
    rebound=post(client,'sources',agent_id=agent['agent']['id'],machine_id=machine['id'],source_id=source['source_id'],adapter_type='simulator')
    assert read(client,'workspace')['machines'][0]['current_run_id'] is None
    raw=event(machine,rebound,2,count='1');raw['stream_id']='replacement-stream';upload(raw)
    assert read(client,'workspace')['machines'][0]['current_run_id']==original_id
    # A new lifecycle can reuse the same native label; cursor evidence must win.
    raw=event(machine,rebound,3,count='0');raw['stream_id']='replacement-stream';raw['source_identity']['generation']='new-job';upload(raw)
    current=read(client,'workspace')['machines'][0]['current_run_id'];assert current and current!=original_id
    raw=event(machine,source,4,count='2');upload(raw)
    assert read(client,'workspace')['machines'][0]['current_run_id']==current
