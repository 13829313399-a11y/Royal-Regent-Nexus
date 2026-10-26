"""Standard-template printing: no lost content, compact rows, original evidence."""
from copy import copy
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services.material_handover_template import print_options

TEMPLATE = Path(__file__).resolve().parents[2] / 'public/templates/material-handover/supplier-delivery-v1.xlsx'


def source(count=10, positions=(8, 9, 17), change=None):
    book = load_workbook(TEMPLATE)
    sheet = book['送货单填写']
    if count != 10:
        delta = count - 10
        merged = [copy(area) for area in sheet.merged_cells.ranges if area.min_row >= 19]
        for area in merged:
            sheet.unmerge_cells(str(area))
        sheet.move_range('A19:M24', rows=delta)
        for area in merged:
            area.shift(row_shift=delta)
            sheet.merge_cells(str(area))
        for row in range(8, 8 + count):
            sheet.cell(row, 12, 'CNY')
            for column in range(1, 14):
                sheet.cell(row, column)._style = copy(sheet.cell(8, column)._style)
            sheet.row_dimensions[row].height = 32
    sheet['B3'] = '测试供应商'
    sheet['B4'] = 'COMPACT-001'
    sheet['G3'] = '2026-10-25'
    sheet['G4'] = '华康测试收货单位'
    sheet['J3'] = '测试联系人'
    sheet['M3'] = '0000-12345678'
    for index, row in enumerate(positions, 1):
        sheet.cell(row, 1, '布料')
        sheet.cell(row, 2, f'F{index:03}')
        sheet.cell(row, 3, '浅蓝色布')
        sheet.cell(row, 6, index)
        sheet.cell(row, 7, '码')
    footer = count + 9
    sheet.cell(footer + 1, 4, '开单测试员')
    sheet.cell(footer + 1, 8, '送货测试员')
    sheet.cell(footer + 1, 11, '收货测试员')
    sheet.cell(footer + 3, 4, '2026-10-26')
    sheet.cell(footer + 3, 8, '外包装完好')
    if change:
        change(book, sheet)
    stream = BytesIO(); book.save(stream); book.close()
    return stream.getvalue()


@pytest.mark.parametrize('count', [10, 30])
def test_only_unused_rows_are_hidden_and_signed_footer_is_preserved(count):
    raw = source(count=count)
    before = bytes(raw)
    options = print_options(raw)
    assert options['sheets'] == ['送货单填写']
    assert options['range'] == f'A1:M{count + 14}'
    assert options['orientation'] == 'landscape' and options['print_mode'] == 'fit_width'
    hidden = options['compact_handover']['hidden_rows']
    assert hidden == [row - 1 for row in range(8, count + 9) if row not in (8, 9, 17)]
    assert raw == before


@pytest.mark.parametrize('cell,value', [('B10','只有编码'), ('C10','只有名称'), ('F10',0), ('K10',0),
    ('M10','需要保留的说明'), ('L10','HKD'), ('F10','=IF(1=1,"",1)'), ('A10','布料')])
def test_incomplete_zero_formula_and_nondefault_content_is_never_hidden(cell, value):
    raw = source(change=lambda book, sheet: setattr(sheet[cell], 'value', value))
    assert 9 not in print_options(raw)['compact_handover']['hidden_rows']


@pytest.mark.parametrize('change', [
    lambda book, sheet: setattr(sheet['B7'], 'value', '供应商自己的列名'),
    lambda book, sheet: setattr(sheet['A19'], 'value', '其他说明'),
    lambda book, sheet: setattr(sheet['N8'], 'value', '额外证据'),
    lambda book, sheet: setattr(sheet['A25'], 'value', '不可忽略的追加内容'),
    lambda book, sheet: book.create_sheet('追加物料'),
])
def test_unfamiliar_or_extended_evidence_keeps_original_layout(change):
    assert print_options(source(change=change)) == {}


@pytest.mark.parametrize('count,positions', [(10,(8,9,17)), (10,tuple(range(8,18))), (25,tuple(range(8,33))), (60,tuple(range(8,68)))])
def test_native_pdf_compacts_signature_and_repeats_headers_without_instruction_pages(tmp_path, count, positions):
    import pdfplumber
    from app.services.document_tools.office_engine import office_executable, uno_python
    from app.services.material_handover_print import source_pdf
    if not office_executable() or not uno_python():
        pytest.skip('Office/UNO renderer not installed')
    raw = source(count=count, positions=positions, change=(
        lambda book, sheet: setattr(sheet.row_dimensions[17], 'hidden', True)
    ) if len(positions) == 3 else None)
    with source_pdf('统一送货单.xlsx', raw, tmp_path, lambda: False) as output:
        assert len(output) == 1 if count == 10 else len(output) > 1
    assert (tmp_path / 'original.xlsx').read_bytes() == raw
    with pdfplumber.open(tmp_path / 'original.pdf') as pdf:
        all_text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
        for index in range(1, len(positions) + 1):
            assert f'F{index:03}' in all_text
        assert '供应商填写说明' not in all_text and 'EX-F001' not in all_text
        assert '开单测试员' in all_text and '外包装完好' in all_text
        assert all('物料编码' in (page.extract_text() or '') for page in pdf.pages)
        assert all(text in (pdf.pages[-1].extract_text() or '') for text in ['签收信息','开单测试员','收货测试员','外包装完好'])
        if len(positions) == 3:
            # Empty rows between materials and before the footer must disappear.
            words = pdf.pages[0].extract_words()
            material = next(word for word in words if word['text'] == 'F003')
            signature = next(word for word in words if word['text'] == '签收信息')
            assert 0 < signature['top'] - material['bottom'] < 45


def test_wrapped_material_name_and_note_expand_without_losing_their_ends(tmp_path):
    import pdfplumber
    from app.services.document_tools.office_engine import office_executable, uno_python
    from app.services.material_handover_print import source_pdf
    if not office_executable() or not uno_python():
        pytest.skip('Office/UNO renderer not installed')
    def change(book, sheet):
        sheet['C8'] = '需要换行的长物料名称' * 4 + 'NAME_END'
        sheet['M8'] = '需要完整保留的送货说明' * 4 + 'NOTE_END'
    raw = source(change=change)
    with source_pdf('长内容.xlsx', raw, tmp_path, lambda: False):
        pass
    with pdfplumber.open(tmp_path / 'original.pdf') as pdf:
        # Imported CJK/Latin glyph bounds differ; the default tolerance may
        # incorrectly interleave adjacent visual lines in extracted text.
        text = ''.join((page.extract_text(y_tolerance=1) or '').replace('\n', '') for page in pdf.pages)
        assert 'NAME_END' in text and 'NOTE_END' in text
        words = pdf.pages[0].extract_words(y_tolerance=1)
        next_row = next(word for word in words if word['text'] == 'F002')
        assert all(next(word for word in words if marker in word['text'])['bottom'] < next_row['top']
                   for marker in ('NAME_END', 'NOTE_END'))
