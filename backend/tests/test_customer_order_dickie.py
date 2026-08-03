from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from fastapi import HTTPException
from openpyxl.styles import Alignment, Border, Font, Side

from app.api.customer_order import _ensure_customer_factory
from app.services import customer_order_dickie as service
from app.services.customer_order_buzzbee import _decrypt_schedule


STANDARD_PAGES = [
    """
    Simba Dickie HK Ltd.
    Release Order Page 1
    Reference: SC700142026/ 1200
    Date of creation: 17.MAR.2026
    Pers. respons. Michelle Wong - Tel: +852 2709 7088
    Mat.No.: 203302028 Mat. EAN:
    4006333075483
    Fire Truck
    Packing: 6PC / 24PC Volume: 2.458 FT3
    Dickie open box
    Quantity Master Contract.No. PO Contract No. Unit Price Delivery Date
    720 PC 500054097/ 10 300481454/ 10 14.60 HKD 11.MAY.2026
    """,
    """
    Reference: SC700142026/ 1200
    ** This order is for Dickie Germany.
    Port of discharge: HAMBURG
    """,
]

MIXED_PAGES = [
    """
    Release Order Page 1
    Reference: SC700143686/ 2000
    Date of creation: 10.APR.2026
    Pers. respons. Joey Tsang - Tel:
    Mat. No.: 2037120232CH Mat. EAN:
    BMW Police/Mini Excavator/Fendt Tractor
    Packing: OPC / 12PC Volume: 1.095 FT3
    open box
    Quantity Master Contract No. PO Contract No. Unit Price Delivery Date
    1,008 PC to be advised 06.JUN.2026
    """,
    "CALENDAR HOLDINGS LLC Port of discharge: HOUSTON",
    """
    Release order Attachment
    Item No. Shipment Master Purchase Delivery
    203722013 336 PC 500054230/ 10 300488338/ 10 700143686/ 2002 17.30 HKD 22.MAY.2026
    203732000 336 PC 500054232/ 10 300488339/ 10 700143686/ 2003 17.00 HKD 22.MAY.2026
    203712034038 336 PC 500053875/ 10 300488337/ 10 700143686/ 2001 16.30 HKD 22.MAY.2026
    """,
]

MULTI_ORDER_PAGES = [
    """
    Release Order Page 1
    Reference: SC700144157/ 100
    Date of creation: 16.APR.2026
    Mat. No.: | 203717007038 Mat. EAN:
    Police Horse Trailer
    Packing: OPC / 6PC
    Quantity Master Contract No. PO Contract No. Unit Price Delivery Date
    120 PC 500047418/ 10 300489811/ 10 57.80 HKD 11.JUL.2026
    """,
    "Release Order Page 2 Port of discharge: MONTREAL",
    """
    Release Order Page: 1
    Reference: SC700144157/ 600
    Date of creation: 16.APR.2026
    Mat. No.: 203307000 Mat. EAN:
    Fire Fighter
    Packing: OPC / 5PC
    Quantity Master Contract No. PO Contract No. Unit Price Delivery Date
    300 PC 500054339/ 10 300492454/ 10 50.80 HKD 11.JUL.2026
    """,
    "Release Order Page: 2 Port of discharge: MONTREAL",
]

LEGACY_ASW_MASTER_ONLY_PAGES = [
    """
    Release Order Page 1
    Reference: SC700130503/ 200
    Date of creation: 27.AUG.2025
    Pers. respons. Olivia Cheng - Tel:
    Mat. No.: 2037120341AS Mat. EAN:
    RC My First NL, 4-asst
    Packing: 4 PC / 12 PC
    open box
    Quantity Master Contract No. PO Contract No. Unit Price Delivery Date
    14,004 PC to be advised 04.JUL.2026
    """,
    "Release Order Page 2 Port of discharge: ROTTERDAM",
    """
    Release order Attachment
    Item No. Shipment Master Purchase Delivery
    990081000007 14,004 PC 300459663/ 50 700130503/ 206 0.60 HKD 19.JUN.2026
    Handling cost of 2037120341AS
    203712034ASW 3,501 PC 300459663/ 60 700130503/ 202 16.10 16.20 HKD 19.JUN.2026
    203712039ASW 3,501 PC 300459663/ 70 700130503/ 203 16 Ao 1ya0 HKD 19.JUN.2026
    203732000ASW 3,501 PC 300459663/ 80 700130503/ 204 1b-s0 yao HKD 19.JUN.2026
    203742019ASW 3,501 PC 300459663/ 90 700130503/ 201 1x, fo 1g40 HKD 19.JUN.2026
    """,
]

LEGACY_ASW_MIXED_CONTRACT_PAGES = [
    """
    Release Order Page 1
    Reference: SC700130518/ 200
    Date of creation: 28.AUG.2025
    Pers. respons. Olivia Cheng - Tel:
    Mat. No.: 2037120342AS Mat. EAN:
    RC My First BE, 4-asst
    Packing: 4 PC / 12 PC
    open box
    Quantity Master Contract No. PO Contract No. Unit Price Delivery Date
    3,204 PC to be advised 04.JUL.2026
    """,
    "Release Order Page 2 Port of discharge: ROTTERDAM",
    """
    Release order Attachment
    Item No. Shipment Master Purchase Delivery
    990081000008 3,204 PC -300463212/ 10 700130518/ 206 0.60 HKD 19.JUN.2026
    Handling cost of 2037120342AS
    203712034AS1 ( BE ) 801 PC 500052799/ 20 300470070/ 10 70013051 8/ 202 16.30 HKD 19.JUN.2026
    203712039ASW. (N L) 801 PC 300470075/ 10 700130518/ 201 17.30 HKD 19.JUN.2026
    203732000ASW 801 PC 300470075/ 20 700130518/ 204 17.00 HKD 19.JUN.2026
    203742019ASW 801 PC 300470075/ 30 70013051 8/ 203 16.10 HKD 19.JUN.2026
    """,
]

LEGACY_TEDI_ATTACHMENT_PAGES = [
    """
    Release Order Page 1
    Reference: SC700136228/ 100
    Pers. respons. Michelle Wong - Tel:
    Mat. No.: 2037120231TE Mat. EAN:
    Lamborghini Police Car/Fendt Tractor
    Packing: OPC / 5PC
    Dickie open box
    Quantity Master Contract No. PO Contract No. Unit Price Delivery Date
    6,340 PC Please refer to attachment
    """,
    "Release Order Page: 2 This order is from Dickie Germany for customer Tedi Germany. Port of discharge: HAMBURG",
    """
    Release order Attachment
    Item No. Shipment Master Purchase Delivery
    203712023TED 2,536 PC 500052620/ 10 300462034/ 10 700136228/ 101 15.80 HKD 04.JUL.2026
    203732000TED 3,804 PC 500052620/ 20 300462040/ 10 700136228/ 103 16.50 HKD 04.JUL.2026
    """,
]


def build_dickie_schedule() -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    order = workbook.create_sheet("接单表")
    review = workbook.create_sheet("正单评审表")
    item = workbook.create_sheet("Iteam表")
    dino = workbook.create_sheet("恐龙蛋口水车Iteam表")
    mo = workbook.create_sheet("MO订单")

    order_headers = {
        "B3": "客出单日期",
        "C3": "主合同",
        "D3": "Reference",
        "E3": "PO.NO",
        "F3": "客名/國家",
        "G3": "產品編號",
        "H3": "產品名称",
        "I3": "數量",
        "J3": "装箱",
        "K3": "单价",
        "L3": "金额",
        "O3": "客要求走货期",
    }
    for coordinate, value in order_headers.items():
        order[coordinate] = value
    order["B875"] = date(2026, 1, 1)
    order["C875"] = "500000001/10"
    order["D875"] = "SC700000001-100"
    for row_number in range(876, 881):
        order[f"D{row_number}"] = "=NA()"
    order["H881"] = "2026年接单"
    order["K881"] = "合计:HK$"
    # Simulate a record misplaced below the subtotal by the previous importer.
    order["B883"] = date(2026, 7, 21)
    order["C883"] = "500055628/10"
    order["D883"] = "SC700149043-500"

    review["B306"] = date(2026, 1, 1)
    review["C306"] = "500000001/10"
    review["D306"] = "SC700000001-100"
    for row_number in range(307, 319):
        review[f"D{row_number}"] = "=NA()"
    # Simulate the old importer writing immediately before the footer.
    review["B319"] = date(2026, 7, 21)
    review["C319"] = "500055628/10"
    review["D319"] = "SC700149043-500"
    review["F320"] = "负责人：罗成灿"

    thin = Side(style="thin", color="FF000000")
    for worksheet, row_number, end_column, font_size in (
        (order, 875, 15, 10),
        (review, 306, 20, 9),
    ):
        worksheet.row_dimensions[row_number].height = 12
        for column_number in range(2, end_column + 1):
            cell = worksheet.cell(row_number, column_number)
            cell.font = Font(
                name="宋体",
                size=font_size,
                color="FF0000FF" if column_number <= 4 else "FF000000",
            )
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(horizontal="center", vertical="center")

    normal_headers = [
        "客出单日期",
        "主合同号",
        "Reference",
        "PO.NO",
        "客名/国家",
        "产品编号",
        "客货号",
        "产品名称",
        "数量",
        "装箱",
        "联系人",
        "包装",
        "合同更改批注",
        "备注",
        "客走货期",
        "inspection date",
        "客验货期",
    ]
    for column, header in enumerate(normal_headers, start=1):
        item.cell(3, column).value = header
    item["A4"] = date(2026, 1, 1)
    item["B4"] = "500000001/10"
    item["C4"] = "SC700000001-100"
    item["D4"] = "300000001/10"
    item["E4"] = "Dickie Germany"
    item["F4"] = "20 330 2028"
    item["G4"] = "20 330 2028"
    item["H4"] = "小消防车"
    item["I4"] = 120
    item["J4"] = "6/24"
    item["K4"] = "Michelle"
    item["L4"] = "Dickie多文盒"
    item["O4"] = date(2026, 6, 1)
    item["Q4"] = date(2026, 5, 25)
    item["A5"] = date(2026, 1, 1)
    item["B5"] = "500054097"
    item["C5"] = "已入系统"
    item["F5"] = "20 330 2028"
    item["H5"] = "小消防车"
    item["I5"] = "=3000"
    item["J5"] = "6/24"

    dino_headers = [
        "客出单日期",
        "主合同号",
        "Reference",
        "PO.NO",
        "客名/国家",
        "产品编号",
        "客货号",
        "产品名称",
        "数量（蛋）",
        "装箱",
        "PDQ数量",
        "联系人",
        "包装",
        "合同更改批注",
        "备注",
        "走货期",
        "inspection date",
        "客验货期",
    ]
    for column, header in enumerate(dino_headers, start=1):
        dino.cell(1, column).value = header
    dino["A2"] = date(2026, 1, 1)
    dino["B2"] = "500055078/10"
    dino["C2"] = "SC700000002-100"
    dino["D2"] = "300000002/10"
    dino["E2"] = "Dickie Germany"
    dino["F2"] = "20 375 3007"
    dino["G2"] = "20 375 3007"
    dino["H2"] = "透明大蛋-警车/消防车/垃圾车"
    dino["I2"] = 60
    dino["J2"] = "0/6"
    dino["L2"] = "Michelle"
    dino["M2"] = "纸箱"
    dino["P2"] = date(2026, 9, 30)
    dino["R2"] = date(2026, 9, 22)
    dino["A3"] = date(2026, 1, 1)
    dino["B3"] = "500055078"
    dino["C3"] = "已入系统"
    dino["F3"] = "20 375 3007"
    dino["I3"] = "=5000"

    mo["A1"] = "客出单日期"
    mo["C1"] = "Item#"
    mo["D1"] = "产品名称"
    mo["E1"] = "数量"
    mo["F1"] = "单价"
    mo["G1"] = "金额"

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_dickie_parser_maps_standard_release_order_and_mixed_allocations():
    standard = service._parse_dickie_ocr_pages(
        STANDARD_PAGES,
        fallback_received_date="2026-07-29",
    )
    assert standard.reference_no == "SC700142026-1200"
    assert standard.master_contract == "500054097/10"
    assert standard.po_no == "300481454/10"
    assert standard.source_date == "2026-03-17"
    assert standard.ship_date == "2026-05-11"
    assert standard.product_no == "203302028"
    assert standard.product_name_en == "Fire Truck"
    assert standard.quantity == 720
    assert standard.unit_price_hkd == Decimal("14.60")
    assert standard.packing == "6/24"
    assert standard.customer_name == "Dickie Germany"
    assert standard.country == "德国"
    assert standard.contact == "Michelle Wong"
    assert standard.packaging == "Dickie open box"

    mixed = service._parse_dickie_ocr_pages(
        MIXED_PAGES,
        fallback_received_date="2026-07-29",
    )
    assert mixed.reference_no == "SC700143686-2000"
    assert mixed.master_contract.splitlines() == [
        "500054230/10",
        "500054232/10",
        "500053875/10",
    ]
    assert mixed.po_no.splitlines() == [
        "300488338/10",
        "300488339/10",
        "300488337/10",
    ]
    assert mixed.quantity == 1008
    assert mixed.unit_price_hkd == Decimal("16.86666666666666666666666667")
    assert mixed.allocations == [
        ("500054230/10", Decimal("336")),
        ("500054232/10", Decimal("336")),
        ("500053875/10", Decimal("336")),
    ]
    assert mixed.ship_date == "2026-06-06"
    assert mixed.customer_name == "Calendar"
    assert mixed.country == "美国"


def test_dickie_parser_handles_legacy_asw_and_tedi_attachment_rows():
    master_only = service._parse_dickie_ocr_pages(
        LEGACY_ASW_MASTER_ONLY_PAGES,
        fallback_received_date="2026-08-03",
    )
    assert master_only.master_contract.splitlines() == [
        "300459663/60",
        "300459663/70",
        "300459663/80",
        "300459663/90",
    ]
    assert master_only.po_no == ""
    assert master_only.quantity == Decimal("14004")
    assert master_only.unit_price_hkd is None
    assert master_only.allocations == [
        ("300459663/60", Decimal("3501")),
        ("300459663/70", Decimal("3501")),
        ("300459663/80", Decimal("3501")),
        ("300459663/90", Decimal("3501")),
    ]

    mixed_contracts = service._parse_dickie_ocr_pages(
        LEGACY_ASW_MIXED_CONTRACT_PAGES,
        fallback_received_date="2026-08-03",
    )
    assert mixed_contracts.master_contract.splitlines() == [
        "500052799/20",
        "300470075/10",
        "300470075/20",
        "300470075/30",
    ]
    assert mixed_contracts.po_no == "300470070/10"
    assert mixed_contracts.quantity == Decimal("3204")
    assert mixed_contracts.unit_price_hkd == Decimal("16.675")

    tedi = service._parse_dickie_ocr_pages(
        LEGACY_TEDI_ATTACHMENT_PAGES,
        fallback_received_date="2026-08-03",
    )
    assert tedi.master_contract.splitlines() == [
        "500052620/10",
        "500052620/20",
    ]
    assert tedi.po_no.splitlines() == [
        "300462034/10",
        "300462040/10",
    ]
    assert tedi.quantity == Decimal("6340")
    assert tedi.unit_price_hkd == Decimal("16.22")
    assert tedi.allocations == [
        ("500052620/10", Decimal("2536")),
        ("500052620/20", Decimal("3804")),
    ]


def test_dickie_parser_splits_combined_pdf_and_accepts_ocr_separator_before_product_no(
    monkeypatch,
):
    monkeypatch.setattr(service, "_extract_pdf_ocr_pages", lambda _content: MULTI_ORDER_PAGES)

    orders = service.parse_dickie_pdf_orders(
        "RoyalRegent(1).pdf",
        b"%PDF-1.7 synthetic",
        fallback_received_date="2026-07-29",
    )

    assert len(orders) == 2
    assert [order.reference_no for order in orders] == [
        "SC700144157-100",
        "SC700144157-600",
    ]
    assert [order.product_no for order in orders] == ["203717007038", "203307000"]
    assert [order.quantity for order in orders] == [Decimal("120"), Decimal("300")]
    assert [order.country for order in orders] == ["加拿大", "加拿大"]


def test_dickie_preview_and_export_write_three_tables_and_deduct_matching_stock(monkeypatch):
    schedule_content = build_dickie_schedule()
    monkeypatch.setattr(service, "_extract_pdf_ocr_pages", lambda _content: STANDARD_PAGES)
    po_files = [("SC700142026-1200.pdf", b"%PDF-1.7 synthetic")]

    preview = service.create_dickie_batch_preview(
        factory_id="huaxing",
        received_date="2026-07-29",
        po_files=po_files,
        schedule_file_name="2026年.Dickie 生产情况.xlsx",
        schedule_content=schedule_content,
    )
    assert preview["summary"] == {"total": 1, "valid": 1, "warning": 0, "blocked": 0}
    row = preview["rows"][0]
    assert row["item_sheet_name"] == "Iteam表"
    assert row["product_no"] == "20 330 2028"
    assert row["product_name_zh"] == "小消防车"
    assert row["units_per_carton"] == "24"
    assert row["carton_count"] == "30"
    assert row["customer_q"] == "2026-05-04"
    assert preview["output_file_name"] == "2026年.Dickie 生产情况.xlsx"

    output, file_name, _ = service.export_dickie_batch_schedule(
        factory_id="huaxing",
        received_date="2026-07-29",
        po_files=po_files,
        schedule_file_name="2026年.Dickie 生产情况.xlsx",
        schedule_content=schedule_content,
    )
    assert file_name == "2026年.Dickie 生产情况.xlsx"
    plain, encrypted = _decrypt_schedule(output)
    assert encrypted is True
    workbook = service.DickieSchedule(plain)

    item_row_number, item_values = next(
        (row_number, values)
        for row_number, values in workbook.read_rows("Iteam表").items()
        if values.get("C") == "SC700142026-1200"
    )
    assert item_values["B"] == "500054097/10"
    assert item_values["D"] == "300481454/10"
    assert item_values["E"] == "Dickie Germany"
    assert item_values["F"] == "20 330 2028"
    assert item_values["H"] == "小消防车"
    assert item_values["I"] == "720"
    assert item_values["J"] == "6/24"
    assert item_values["K"] == "Michelle Wong"
    assert item_values["L"] == "Dickie open box"
    assert _format_serial(item_values["O"]) == "2026-05-11"
    assert _format_serial(item_values["Q"]) == "2026-05-04"

    system_row_number = next(
        row_number
        for row_number, values in workbook.read_rows("Iteam表").items()
        if values.get("C") == "已入系统"
        and values.get("F") == "20 330 2028"
    )
    assert workbook.read_cell_formula("Iteam表", system_row_number, "I") == "3000-720"

    order_row_number = next(
        row_number
        for row_number, values in workbook.read_rows("接单表").items()
        if values.get("D") == "SC700142026-1200"
    )
    assert order_row_number == 876
    assert workbook.read_cell_formula("接单表", order_row_number, "E") == (
        f"VLOOKUP(D{order_row_number},'Iteam表'!C:D,2,0)"
    )
    assert workbook.read_cell_formula("接单表", order_row_number, "L") == (
        f"I{order_row_number}*K{order_row_number}"
    )

    review_row_number = next(
        row_number
        for row_number, values in workbook.read_rows("正单评审表").items()
        if values.get("D") == "SC700142026-1200"
    )
    assert review_row_number == 307
    assert workbook.read_rows("正单评审表")[review_row_number]["S"] == "14.6"
    assert workbook.read_cell_formula("正单评审表", review_row_number, "T") == (
        f"I{review_row_number}*S{review_row_number}"
    )
    assert item_row_number < system_row_number

    rendered = openpyxl.load_workbook(BytesIO(plain), data_only=False)
    for column_number in range(2, 16):
        inserted = rendered["接单表"].cell(order_row_number, column_number)
        reference = rendered["接单表"].cell(order_row_number - 1, column_number)
        assert inserted._style.fontId == reference._style.fontId
        assert inserted._style.borderId == reference._style.borderId
        assert inserted._style.alignmentId == reference._style.alignmentId
        assert inserted._style.numFmtId == reference._style.numFmtId
    for column_number in range(2, 21):
        inserted = rendered["正单评审表"].cell(review_row_number, column_number)
        reference = rendered["正单评审表"].cell(review_row_number - 1, column_number)
        assert inserted._style.fontId == reference._style.fontId
        assert inserted._style.borderId == reference._style.borderId
        assert inserted._style.alignmentId == reference._style.alignmentId
        assert inserted._style.numFmtId == reference._style.numFmtId
    assert rendered["Iteam表"][f"F{item_row_number}"].fill.fgColor.rgb == "FFFFFF00"
    assert rendered["接单表"][f"D{order_row_number}"].fill.fgColor.rgb == "FFFFFF00"
    assert rendered["正单评审表"][f"D{review_row_number}"].fill.fgColor.rgb == "FFFFFF00"
    assert rendered["Iteam表"][f"F{system_row_number}"].fill.fgColor.rgb != "FFFFFF00"
    rendered.close()


def _format_serial(value: str) -> str:
    return service._format_iso_date(Decimal(value))


def test_dickie_preview_blocks_reference_already_in_current_schedule(monkeypatch):
    workbook = openpyxl.load_workbook(BytesIO(build_dickie_schedule()))
    workbook["Iteam表"]["C4"] = "SC700142026-1200"
    schedule_output = BytesIO()
    workbook.save(schedule_output)
    workbook.close()
    monkeypatch.setattr(service, "_extract_pdf_ocr_pages", lambda _content: STANDARD_PAGES)

    preview = service.create_dickie_batch_preview(
        factory_id="huaxing",
        received_date="2026-07-29",
        po_files=[("SC700142026-1200.pdf", b"%PDF-1.7 synthetic")],
        schedule_file_name="2026年.Dickie 生产情况.xlsx",
        schedule_content=schedule_output.getvalue(),
    )

    assert preview["summary"] == {"total": 1, "valid": 0, "warning": 0, "blocked": 1}
    duplicate_issue = next(
        issue
        for issue in preview["rows"][0]["issues"]
        if issue["code"] == "duplicate_reference"
    )
    assert duplicate_issue["severity"] == "blocked"
    assert "Iteam表第 4 行" in duplicate_issue["message"]


def test_customer_factory_mapping_rejects_cross_factory_imports():
    _ensure_customer_factory("buzzbee", "huaxing")
    _ensure_customer_factory("dickie", "huaxing")
    with pytest.raises(HTTPException, match="只属于华兴厂区"):
        _ensure_customer_factory("dickie", "huakang-a")
