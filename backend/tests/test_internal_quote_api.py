import importlib
import hashlib
import sys
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook


TEST_TMP_DIR = Path(__file__).resolve().parents[1] / ".pytest-tmp"
BACKEND_DIR = Path(__file__).resolve().parents[1]
ADMIN_TEST_PASSWORD = "AdminSeed123!"

if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))


def make_client(monkeypatch) -> TestClient:
    TEST_TMP_DIR.mkdir(exist_ok=True)
    database_url = f"sqlite:///{TEST_TMP_DIR / f'internal_quote_{uuid4().hex}.db'}"
    monkeypatch.setenv("DATABASE_URL", database_url)
    monkeypatch.setenv("SEED_ADMIN_PASSWORD", ADMIN_TEST_PASSWORD)

    for module_name in list(sys.modules):
        if module_name == "app" or module_name.startswith("app."):
            del sys.modules[module_name]

    main = importlib.import_module("app.main")
    return TestClient(main.app)


def create_test_user(username: str, role_id: str, factory_id: str = "huaxing") -> str:
    db_module = importlib.import_module("app.db")
    auth_models = importlib.import_module("app.models.auth")
    auth_service = importlib.import_module("app.services.auth")
    user_id = f"user-{username}"

    with db_module.SessionLocal() as db:
        salt, password_hash = auth_service.make_password_hash("123456")
        db.add(
            auth_models.AuthUser(
                id=user_id,
                username=username,
                display_name=username,
                password_salt=salt,
                password_hash=password_hash,
                status="active",
                force_password_change=0,
                created_at=auth_service.now_text(),
                updated_at=auth_service.now_text(),
            )
        )
        db.flush()
        db.add(
            auth_models.AuthUserRole(
                id=f"{user_id}:{role_id}:{factory_id}:sales-business",
                user_id=user_id,
                role_id=role_id,
                factory_id=factory_id,
                department="sales-business",
            )
        )
        db.commit()
    return user_id


def login(client: TestClient, username: str):
    response = client.post("/api/auth/login", json={"username": username, "password": "123456"})
    assert response.status_code == 200
    return response.json()


def create_quote(client: TestClient, quote_no: str = "47765A"):
    return client.post(
        "/api/internal-quotes",
        json={
            "factory_id": "huaxing",
            "workshop_code": "huaxing-workshop",
            "quote_no": quote_no,
            "product_name": "八部门协同报价产品",
            "customer": "Target",
            "qty": 10000,
            "version_label": "V1",
        },
    )


def section_payload(index: int = 1):
    return {
        "currency": "HKD",
        "loss_pct": 5,
        "rows": [
            {
                "id": f"line-{index}",
                "category": "材料",
                "item_name": f"成本项 {index}",
                "specification": "标准",
                "quantity": 2,
                "unit_price_hkd": 10,
                "amount_hkd": 999999,
                "note": "由服务器复算",
            }
        ],
    }


def test_legacy_superadmin_can_read_export_history_when_role_mapping_predates_permission(monkeypatch):
    monkeypatch.setenv("AUTHZ_MODE", "legacy")
    monkeypatch.setenv("AUTHZ_WRITES_ENABLED", "false")
    with make_client(monkeypatch) as client:
        db_module = importlib.import_module("app.db")
        auth_models = importlib.import_module("app.models.auth")
        with db_module.SessionLocal() as db:
            permission = db.query(auth_models.AuthPermission).filter_by(
                code="internal_pricing:export",
            ).one()
            mapping = db.query(auth_models.AuthRolePermission).filter_by(
                role_id="admin",
                permission_id=permission.id,
            ).one()
            db.delete(mapping)
            db.commit()

        login_response = client.post(
            "/api/auth/login",
            json={"username": "admin", "password": ADMIN_TEST_PASSWORD},
        )
        assert login_response.status_code == 200
        profile = login_response.json()
        assert "internal_pricing:export" in profile["permissions"]
        assert "internal_pricing:export" not in profile["grants"][0]["permissions"]

        created = create_quote(client, "Q-LEGACY-ADMIN")
        assert created.status_code == 201
        export_history = client.get(f"/api/internal-quotes/{created.json()['id']}/exports")
        assert export_history.status_code == 200
        assert export_history.json() == []


def test_workshop_quote_runs_eight_section_approval_and_controlled_export(monkeypatch):
    with make_client(monkeypatch) as client:
        create_test_user("quote_owner", "sales_customer_owner")
        create_test_user("quote_supervisor", "sales_customer_supervisor")
        owner_profile = login(client, "quote_owner")
        assert {
            "internal_pricing:create",
            "internal_pricing:edit",
            "internal_pricing:export",
        }.issubset(owner_profile["permissions"])

        workshops = client.get("/api/internal-quotes/workshops?factory_id=huaxing")
        assert workshops.status_code == 200
        assert workshops.json() == [{"code": "huaxing-workshop", "name": "华兴"}]

        invalid_workshop = client.post(
            "/api/internal-quotes",
            json={
                "factory_id": "huaxing",
                "workshop_code": "unknown",
                "quote_no": "INVALID",
                "product_name": "错误车间",
                "customer": "Target",
                "qty": 1,
                "version_label": "V1",
            },
        )
        assert invalid_workshop.status_code == 400

        created_response = create_quote(client)
        assert created_response.status_code == 201
        detail = created_response.json()
        quote_id = detail["id"]
        assert detail["workshop_code"] == "huaxing-workshop"
        assert detail["workshop_name"] == "华兴"
        assert len(detail["sections"]) == 8
        assert {section["department"] for section in detail["sections"]} == {
            "sales",
            "engineering",
            "electronic",
            "molding",
            "painting",
            "slush",
            "sewing",
            "assembly",
        }

        huaxing_list = client.get(
            "/api/internal-quotes?factory_id=huaxing&workshop_code=huaxing-workshop"
        )
        assert [quote["id"] for quote in huaxing_list.json()] == [quote_id]

        for index, section in enumerate(detail["sections"], start=1):
            update = client.put(
                f"/api/internal-quotes/{quote_id}/sections/{section['department']}",
                json={
                    "revision": section["revision"],
                    "payload": section_payload(index),
                    "submit": True,
                },
            )
            assert update.status_code == 200
            detail = update.json()

        sales = next(section for section in detail["sections"] if section["department"] == "sales")
        assert sales["status"] == "pending_review"
        assert sales["payload"]["rows"][0]["amount_hkd"] == 20
        assert sales["calculation"]["formula_version"] == "department-formulas-v1"
        assert sales["calculation"]["subtotal_hkd"] == 20
        assert sales["calculation"]["loss_amount_hkd"] == 1
        assert sales["calculation"]["total_hkd"] == 21
        assert sales["calculation"]["total_rmb"] == 17.85
        assert sales["calculation"]["total_usd"] == 2.69
        assert sales["calculation"]["line_breakdown"][0]["amount_hkd"] == 20
        assert sales["payload"]["reference_snapshot"]["version"] == "rr2-2026-v1"
        assert sales["payload"]["reference_snapshot"]["fx"] == {
            "rmb_hkd": 0.85,
            "hkd_usd": 7.8,
        }
        assert detail["total_hkd"] == 168

        locked_update = client.put(
            f"/api/internal-quotes/{quote_id}/sections/sales",
            json={
                "revision": sales["revision"],
                "payload": section_payload(99),
                "submit": True,
            },
        )
        assert locked_update.status_code == 409
        assert locked_update.json()["detail"] == "该分段正在审核中，不能修改"

        stale_update = client.put(
            f"/api/internal-quotes/{quote_id}/sections/engineering",
            json={
                "revision": 1,
                "payload": section_payload(99),
                "submit": False,
            },
        )
        assert stale_update.status_code == 409
        assert stale_update.json()["detail"]["message"] == "报价分段已被其他人更新，请刷新后重试"

        supervisor_profile = login(client, "quote_supervisor")
        assert "internal_pricing:review" in supervisor_profile["permissions"]
        for section in detail["sections"]:
            review = client.post(
                f"/api/internal-quotes/{quote_id}/sections/{section['department']}/review",
                json={"action": "approve", "comment": "主管复核通过"},
            )
            assert review.status_code == 200
            detail = review.json()

        assert detail["status"] == "fully_approved"
        assert detail["approved_count"] == 8

        export = client.get(f"/api/internal-quotes/{quote_id}/export")
        assert export.status_code == 200
        assert export.content[:2] == b"PK"
        assert export.headers["x-content-sha256"] == hashlib.sha256(export.content).hexdigest()
        assert export.headers["x-export-id"]
        workbook = load_workbook(BytesIO(export.content), read_only=True, data_only=True)
        assert workbook.sheetnames == ["内部报价明细"]
        assert workbook["内部报价明细"]["A1"].value == "47765A 八部门协同报价产品 内部报价明细"

        exported_detail = client.get(f"/api/internal-quotes/{quote_id}").json()
        assert exported_detail["status"] == "exported"
        assert exported_detail["audit_logs"][0]["action"] == "export"

        export_history = client.get(f"/api/internal-quotes/{quote_id}/exports")
        assert export_history.status_code == 200
        assert len(export_history.json()) == 1
        retained = export_history.json()[0]
        assert retained["id"] == export.headers["x-export-id"]
        assert retained["status"] == "current"
        assert retained["sha256"] == hashlib.sha256(export.content).hexdigest()
        assert retained["section_revisions"] == {
            section["department"]: section["revision"] for section in exported_detail["sections"]
        }

        retained_download = client.get(
            f"/api/internal-quotes/{quote_id}/exports/{retained['id']}/download"
        )
        assert retained_download.status_code == 200
        assert retained_download.content == export.content
        assert retained_download.headers["x-export-status"] == "current"

        reopened = client.post(
            f"/api/internal-quotes/{quote_id}/sections/sales/review",
            json={"action": "reopen", "comment": "客户规格发生变化"},
        )
        assert reopened.status_code == 200
        assert reopened.json()["status"] == "reopened"
        superseded = client.get(f"/api/internal-quotes/{quote_id}/exports").json()[0]
        assert superseded["status"] == "superseded"
        assert superseded["superseded_at"]


def test_scope_and_maker_checker_rules_are_enforced_by_backend(monkeypatch):
    with make_client(monkeypatch) as client:
        create_test_user("self_reviewer", "sales_customer_supervisor")
        create_test_user("wrong_factory", "sales_customer_owner", factory_id="huadeng")
        login(client, "self_reviewer")

        created = create_quote(client, quote_no="SELF-CHECK")
        assert created.status_code == 201
        detail = created.json()
        sales = next(section for section in detail["sections"] if section["department"] == "sales")
        submitted = client.put(
            f"/api/internal-quotes/{detail['id']}/sections/sales",
            json={
                "revision": sales["revision"],
                "payload": section_payload(),
                "submit": True,
            },
        )
        assert submitted.status_code == 200

        self_review = client.post(
            f"/api/internal-quotes/{detail['id']}/sections/sales/review",
            json={"action": "approve", "comment": "自己审核"},
        )
        assert self_review.status_code == 409
        assert self_review.json()["detail"] == "提交人与审核人不能为同一账号"

        login(client, "wrong_factory")
        denied = client.get("/api/internal-quotes?factory_id=huaxing")
        assert denied.status_code == 403


def test_excel_preview_confirm_and_attachment_lifecycle(monkeypatch):
    with make_client(monkeypatch) as client:
        create_test_user("artifact_owner", "sales_customer_owner")
        login(client, "artifact_owner")
        created = create_quote(client, quote_no="IMPORT-001")
        assert created.status_code == 201
        quote = created.json()
        quote_id = quote["id"]
        electronic = next(
            section for section in quote["sections"] if section["department"] == "electronic"
        )

        source = Workbook()
        sheet = source.active
        sheet.title = "电子报价"
        sheet.append(["零件名称", "规格", "用量", "单价", "备注"])
        sheet.append(["IC", "A1", 2, 10, "主控"])
        sheet.append(["邦定成本", 1.5])
        output = BytesIO()
        source.save(output)
        source.close()

        preview_response = client.post(
            f"/api/internal-quotes/{quote_id}/imports/electronic/preview",
            files={
                "file": (
                    "电子报价.xlsx",
                    output.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert preview_response.status_code == 200
        preview = preview_response.json()
        assert preview["target_department"] == "electronic"
        assert preview["status"] == "previewed"
        assert len(preview["rows"]) == 1
        assert preview["rows"][0]["fields"]["unit_price_rmb"] == 10
        assert preview["parameters"]["bonding_cost_rmb"] == 1.5
        assert preview["source_sha256"] == hashlib.sha256(output.getvalue()).hexdigest()

        unchanged = client.get(f"/api/internal-quotes/{quote_id}").json()
        unchanged_electronic = next(
            section for section in unchanged["sections"] if section["department"] == "electronic"
        )
        assert unchanged_electronic["revision"] == electronic["revision"]
        assert unchanged_electronic["payload"]["rows"] == electronic["payload"]["rows"]

        confirmed_response = client.post(
            f"/api/internal-quotes/{quote_id}/imports/{preview['batch_id']}/confirm",
            json={"revision": electronic["revision"], "mode": "replace"},
        )
        assert confirmed_response.status_code == 200
        confirmed = next(
            section
            for section in confirmed_response.json()["sections"]
            if section["department"] == "electronic"
        )
        assert confirmed["revision"] == electronic["revision"] + 1
        assert confirmed["payload"]["rows"][0]["item_name"] == "IC"
        assert confirmed["payload"]["parameters"]["bonding_cost_rmb"] == 1.5

        duplicate = client.post(
            f"/api/internal-quotes/{quote_id}/imports/{preview['batch_id']}/confirm",
            json={"revision": confirmed["revision"], "mode": "append"},
        )
        assert duplicate.status_code == 409

        imports = client.get(f"/api/internal-quotes/{quote_id}/imports")
        assert imports.status_code == 200
        assert imports.json()[0]["status"] == "confirmed"

        attachment_content = b"quote reference document"
        uploaded_response = client.post(
            f"/api/internal-quotes/{quote_id}/attachments",
            data={"department": "electronic"},
            files={"file": ("reference.pdf", attachment_content, "application/pdf")},
        )
        assert uploaded_response.status_code == 200
        uploaded = uploaded_response.json()
        assert uploaded["department"] == "electronic"
        assert uploaded["sha256"] == hashlib.sha256(attachment_content).hexdigest()

        attachment_list = client.get(f"/api/internal-quotes/{quote_id}/attachments")
        assert attachment_list.status_code == 200
        assert attachment_list.json() == [uploaded]

        downloaded = client.get(
            f"/api/internal-quotes/{quote_id}/attachments/{uploaded['id']}/download"
        )
        assert downloaded.status_code == 200
        assert downloaded.content == attachment_content
        assert downloaded.headers["x-content-sha256"] == uploaded["sha256"]
