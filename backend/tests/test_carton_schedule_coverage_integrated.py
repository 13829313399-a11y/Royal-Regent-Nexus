"""Independent coverage oracles through real XLSX/API transactions, synthetic DBs only."""
from datetime import datetime
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook

from test_molding_sample_api import make_client, login_as
from test_carton_procurement_api import _ensure_dickie_customer, _order_payload

BASE = '/api/carton-procurement'
COVERAGE = BASE + '/schedule-coverage'
SHEET = 'ITEM业务排期'
HEADERS = ['订单类型', 'Contract No.', 'SO#/Reference', 'P/O#:', '客名', '产品编号',
           '产品中文名称', '數量', '装箱', '验货日期', '客要求走货期']


@pytest.fixture(autouse=True)
def isolated(monkeypatch, tmp_path):
    import test_molding_sample_api
    monkeypatch.setattr(test_molding_sample_api, 'TEST_TMP_DIR', tmp_path)


def prepare(client, monkeypatch):
    login_as(client, 'admin')
    _ensure_dickie_customer(client)
    login_as(client, 'admin')
    from app.services import carton_procurement as core
    monkeypatch.setattr(core, 'business_now', lambda: datetime(2026, 10, 20, 12))
    monkeypatch.setattr(core, 'business_today', lambda: '2026-10-20')


def row(contract, quantity=360, due='2026-09-17'):
    return ['正单', contract, 'SO-' + contract, 'PO-' + contract, 'Dickie', 'ITEM-' + contract,
            '覆盖核对产品', quantity, 6, '2026-09-10', due]


def workbook(rows):
    book = Workbook()
    book.active.title = SHEET
    book.active.append(HEADERS)
    for values in rows:
        book.active.append(values)
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def scope(mode, day, **extra):
    values = dict(scope_mode=mode, as_of_date=day)
    if mode == 'FULL':
        values.update(coverage_start='2026-09-01', coverage_end='2026-09-30', full_scope_confirmed='true')
    return {**values, **extra}


def upload(client, name, rows=None, *, content=None, declared=None, expected=201):
    response = client.post(BASE + '/weekly-imports', params={'factory_id': 'huaxing'},
        data={'customer_code': 'DICKIE', **(declared or {})},
        files={'file': (name, content if content is not None else workbook(rows))})
    assert response.status_code == expected, response.text
    return response.json()


def workspace(client, status='ALL', factory='huaxing'):
    response = client.get(COVERAGE + '/workspace', params=dict(factory_id=factory, status=status, limit=100))
    assert response.status_code == 200, response.text
    return response.json()


def missing(client, contract):
    return next(item for item in workspace(client)['missing']['items'] if item['contract_no'] == contract)


def detail(client, case_id):
    response = client.get(COVERAGE + f'/missing/{case_id}', params={'factory_id': 'huaxing'})
    assert response.status_code == 200, response.text
    return response.json()


def review_body(case, decision, **extra):
    return dict(factory_id='huaxing', request_id=uuid4().hex, expected_revision=case['revision'],
        source_fingerprint=case['source_fingerprint'], decision=decision,
        reason='业务及仓库核实原需求和本次排期遗漏', **extra)


def review(client, case, decision, *, expected=200):
    body = review_body(case, decision)
    response = client.post(COVERAGE + f"/missing/{case['id']}/review", json=body)
    assert response.status_code == expected, response.text
    return response.json(), body


def undo(client, batch):
    response = client.post(BASE + f"/imports/{batch['id']}/undo", json=dict(
        factory_id='huaxing', reason='本次排期文件有误，整批撤销核实'))
    assert response.status_code == 200, response.text
    return response.json()


def source(batch, contract):
    value = next(item for item in batch['parse_summary']['rows'] if item.get('contract_no') == contract)
    return dict(batch_id=batch['id'], source_sheet=value['source_sheet'], source_row=value['source_row'])


def order_body(batch, contract):
    value = next(item for item in batch['parse_summary']['rows'] if item.get('contract_no') == contract)
    return {**_order_payload(), 'request_id': uuid4().hex, 'contract_no': contract,
        'item_no': value['item_no'], 'customer_po': value['source_reference'],
        'product_order_quantity': value['quantity'], 'customer_due_date': value['customer_due_date'],
        'schedule_source': source(batch, contract)}


def mark(client, batch, contract, marked=True, *, expected=200, reason='核实原纸箱采购需求仍然有效'):
    response = client.post(BASE + '/schedule-order-marks', json=dict(
        factory_id='huaxing', **source(batch, contract), marked=marked, reason=reason))
    assert response.status_code == expected, response.text
    return response.json()


def business_snapshot():
    """Coverage may append its own audit, never touch any demand or stock fact."""
    from sqlalchemy import select
    from app.db import SessionLocal
    from app.models.carton_procurement import (CartonOrder, CartonOrderLine, CartonInventoryMovement,
        CartonReceipt, CartonReceiptLine, CartonPurchaseOrderIssue, CartonAuditEvent)
    result = {}
    with SessionLocal() as db:
        for model in (CartonOrder, CartonOrderLine, CartonInventoryMovement,
                      CartonReceipt, CartonReceiptLine, CartonPurchaseOrderIssue):
            table = model.__table__
            result[table.name] = [dict(item) for item in db.execute(select(table).order_by(table.c.id)).mappings()]
        result['marks'] = [event.detail_json for event in db.scalars(select(CartonAuditEvent).where(
            CartonAuditEvent.event_type == 'SCHEDULE_ORDER_MARK_CHANGED').order_by(CartonAuditEvent.sequence))]
    return result


def test_full_partial_reappearance_and_withdrawal_keep_demand_stock_and_manual_evidence(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client, monkeypatch)
        values = [row(key) for key in ('ORDERED1', 'MARKED2', 'NEW3', 'STAYS4')]
        baseline = upload(client, 'full-baseline.xlsx', values, declared=scope('FULL', '2026-10-10'))
        created = client.post(BASE + '/orders', json=order_body(baseline, 'ORDERED1'))
        assert created.status_code == 201, created.text
        mark(client, baseline, 'MARKED2')
        before = business_snapshot()

        upload(client, 'partial.xlsx', [values[-1]], declared=scope('INCREMENTAL', '2026-10-11'))
        assert workspace(client, 'PENDING')['missing']['total'] == 0
        full_content = workbook([values[-1]])
        full = upload(client, 'full-missing.xlsx', content=full_content, declared=scope('FULL', '2026-10-12'))
        cases = workspace(client, 'PENDING')['missing']['items']
        assert {item['contract_no'] for item in cases} == {'ORDERED1', 'MARKED2', 'NEW3'}
        assert all(item['status'] == 'PENDING' for item in cases)
        assert business_snapshot() == before

        blocked_body = order_body(baseline, 'NEW3')
        assert client.post(BASE + '/orders', json=blocked_body).status_code == 409
        mark(client, baseline, 'NEW3', expected=409)
        mark(client, baseline, 'NEW3', marked=False)
        assert business_snapshot() == before
        kept, body = review(client, missing(client, 'NEW3'), 'KEEP_REQUIRED')
        assert kept['can_operate_source']
        replay = client.post(COVERAGE + f"/missing/{kept['id']}/review", json=body)
        assert replay.status_code == 200, replay.text
        assert detail(client, kept['id'])['revision'] == kept['revision']
        no_longer, _ = review(client, missing(client, 'MARKED2'), 'NO_LONGER_REQUIRED')
        assert not no_longer['can_operate_source']
        mark(client, baseline, 'MARKED2', expected=409)
        assert business_snapshot() == before
        allowed = client.post(BASE + '/orders', json=blocked_body)
        assert allowed.status_code == 201, allowed.text
        after_create = business_snapshot()

        returned_values = [list(value) for value in values]
        returned_values[1][7] = 720
        returned = upload(client, 'full-returned.xlsx', returned_values, declared=scope('FULL', '2026-10-13'))
        assert detail(client, cases[0]['id'])['status'] == 'RETURNED'
        assert workspace(client, 'PENDING')['missing']['total'] == 0
        mark(client, returned, 'MARKED2', reason='', expected=422)
        assert business_snapshot() == after_create
        undo(client, returned)
        assert missing(client, 'ORDERED1')['status'] == 'PENDING'
        assert missing(client, 'MARKED2')['status'] == 'NO_LONGER_REQUIRED'
        assert business_snapshot() == after_create
        undo(client, full)
        assert workspace(client, 'PENDING')['missing']['total'] == 0
        assert business_snapshot() == after_create

        again = upload(client, 'full-missing.xlsx', content=full_content, declared=scope('FULL', '2026-10-12'))
        assert again['id'] != full['id'] and not again['duplicate']
        assert again['parse_summary']['reimported_from_batch_id'] == full['id']
        fresh = workspace(client, 'PENDING')['missing']['items']
        assert {item['contract_no'] for item in fresh} == {'ORDERED1', 'MARKED2', 'NEW3'}
        assert {item['id'] for item in fresh}.isdisjoint({item['id'] for item in cases})
        assert business_snapshot() == after_create


def test_scope_chronology_legacy_replay_unknown_dates_and_full_parse_gaps(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client, monkeypatch)
        legacy_content = workbook([row('INSIDE1'), row('UNKNOWN2', due='待业务核实'), row('OUTSIDE3', due='2026-10-02')])
        legacy = upload(client, 'legacy.xlsx', content=legacy_content)
        before = business_snapshot()
        full = upload(client, 'full.xlsx', [row('STAYS4')], declared=scope('FULL', '2026-10-12'))
        assert {item['contract_no'] for item in workspace(client, 'PENDING')['missing']['items']} == {'INSIDE1'}
        replay = upload(client, 'same-legacy.xlsx', content=legacy_content)
        assert replay['duplicate'] and replay['id'] == legacy['id']
        upload(client, 'new-undeclared.xlsx', [row('INSIDE1', 400)], expected=409)
        upload(client, 'late-old.xlsx', [row('INSIDE1', 500)], declared=scope('INCREMENTAL', '2026-10-11'), expected=409)
        same_day = upload(client, 'same-day-revision.xlsx', [row('INSIDE1', 500)], declared=scope('INCREMENTAL', '2026-10-12'))
        assert same_day['id'] != full['id']
        assert workspace(client, 'PENDING')['missing']['total'] == 0
        damaged = row('BROKEN5'); damaged[1] = ''
        upload(client, 'full-damaged.xlsx', [row('STAYS4'), damaged], declared=scope('FULL', '2026-10-13'), expected=422)
        book = load_workbook(BytesIO(workbook([row('STAYS4')])))
        extra = book.create_sheet('未识别业务明细'); extra.append(['未知合同', '数量']); extra.append(['INSIDE1', 500])
        stream = BytesIO(); book.save(stream)
        upload(client, 'full-unknown-sheet.xlsx', content=stream.getvalue(), declared=scope('FULL', '2026-10-13'), expected=422)
        assert business_snapshot() == before


def test_withdrawing_referenced_demand_cannot_release_a_still_effective_full_absence(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client, monkeypatch)
        first = upload(client, 'a1.xlsx', [row('ANCESTOR1', 100), row('STAYS2')], declared=scope('FULL', '2026-10-10'))
        revised = upload(client, 'a2.xlsx', [row('ANCESTOR1', 200)], declared=scope('INCREMENTAL', '2026-10-11'))
        full = upload(client, 'b-missing.xlsx', [row('STAYS2')], declared=scope('FULL', '2026-10-12'))
        case = missing(client, 'ANCESTOR1')
        assert {item['quantity'] for item in case['source_rows']} == {200}
        frozen_rows = case['source_rows']
        stale_decision = review_body(case, 'KEEP_REQUIRED')
        before = business_snapshot()
        undo(client, revised)
        assert detail(client, case['id'])['source_rows'] == frozen_rows
        attempted = client.post(COVERAGE + f"/missing/{case['id']}/review", json=stale_decision)
        assert attempted.status_code == 409, attempted.text
        assert client.post(BASE + '/orders', json=order_body(first, 'ANCESTOR1')).status_code == 409
        mark(client, first, 'ANCESTOR1', expected=409)
        assert business_snapshot() == before
        undo(client, full)
        assert workspace(client, 'PENDING')['missing']['total'] == 0
        assert business_snapshot() == before
        restored = client.post(BASE + '/orders', json=order_body(first, 'ANCESTOR1'))
        assert restored.status_code == 201, restored.text


def test_missing_review_preserves_factory_denies_and_refreshes_actor_after_factory_lock(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client, monkeypatch)
        upload(client, 'base.xlsx', [row('SECURE1'), row('STAYS2')], declared=scope('FULL', '2026-10-10'))
        upload(client, 'missing.xlsx', [row('STAYS2')], declared=scope('FULL', '2026-10-11'))
        case = missing(client, 'SECURE1')
        before = business_snapshot()
        foreign = client.get(COVERAGE + f"/missing/{case['id']}", params={'factory_id': 'huadeng'})
        assert foreign.status_code == 404, foreign.text
        from app.db import SessionLocal
        from app.models.auth import AuthUser, AuthPermission, AuthUserPermissionOverride
        from sqlalchemy import select
        with SessionLocal() as db:
            actor = db.scalar(select(AuthUser).where(AuthUser.username == 'warehouse_keeper'))
            permission = db.scalar(select(AuthPermission).where(AuthPermission.code == 'carton_procurement:import'))
            deny_id = 'coverage-deny-' + uuid4().hex
            db.add(AuthUserPermissionOverride(id=deny_id, user_id=actor.id, permission_id=permission.id,
                effect='deny', factory_id='huaxing', department='*', status='active', reason='独立验证显式拒绝仍然优先'))
            db.commit()
        login_as(client, 'warehouse_keeper')
        denied = client.post(COVERAGE + f"/missing/{case['id']}/review", json=review_body(case, 'KEEP_REQUIRED'))
        assert denied.status_code == 403, denied.text
        with SessionLocal() as db:
            db.get(AuthUserPermissionOverride, deny_id).status = 'revoked'; db.commit()
        login_as(client, 'warehouse_keeper')
        from app.services import carton_procurement as core
        original_lock = core._lock_receipt_factory
        touched = []
        def locked_then_disabled(db, factory):
            original_lock(db, factory)
            user = db.scalar(select(AuthUser).where(AuthUser.username == 'warehouse_keeper'))
            user.status = 'disabled'; db.flush(); touched.append(True)
        with monkeypatch.context() as mutation:
            mutation.setattr(core, '_lock_receipt_factory', locked_then_disabled)
            refused = client.post(COVERAGE + f"/missing/{case['id']}/review", json=review_body(case, 'KEEP_REQUIRED'))
        assert touched and refused.status_code == 403, refused.text
        login_as(client, 'admin')
        unchanged = detail(client, case['id'])
        assert unchanged['revision'] == case['revision'] and unchanged['status'] == 'PENDING'
        assert business_snapshot() == before


def test_returned_nonformal_and_zero_rows_block_old_sources_but_latest_duplicate_coordinates_work(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client, monkeypatch)
        baseline = upload(client, 'base.xlsx', [row('REAPPEAR1'), row('STAYS2')], declared=scope('FULL', '2026-10-10'))
        upload(client, 'missing.xlsx', [row('STAYS2')], declared=scope('FULL', '2026-10-11'))
        case = missing(client, 'REAPPEAR1')
        before = business_snapshot()
        nonformal = row('REAPPEAR1'); nonformal[0] = '备料单'
        upload(client, 'nonformal.xlsx', [nonformal], declared=scope('INCREMENTAL', '2026-10-12'))
        returned = detail(client, case['id'])
        assert returned['status'] == 'RETURNED' and not returned['can_operate_source']
        assert client.post(BASE + '/orders', json=order_body(baseline, 'REAPPEAR1')).status_code == 409
        mark(client, baseline, 'REAPPEAR1', expected=409)
        upload(client, 'zero.xlsx', [row('REAPPEAR1', 0)], declared=scope('INCREMENTAL', '2026-10-13'))
        assert not detail(client, case['id'])['can_operate_source']
        assert client.post(BASE + '/orders', json=order_body(baseline, 'REAPPEAR1')).status_code == 409

        duplicate = upload(client, 'duplicates.xlsx', [row('REAPPEAR1'), row('REAPPEAR1', 720)],
            declared=scope('INCREMENTAL', '2026-10-14'))
        old_view = client.get(BASE + f"/imports/{baseline['id']}", params={'factory_id': 'huaxing'})
        assert old_view.status_code == 200, old_view.text
        old_source = next(item for item in old_view.json()['parse_summary']['rows'] if item['contract_no'] == 'REAPPEAR1')
        assert old_source['can_operate_source'] is False
        current = client.get(BASE + f"/imports/{duplicate['id']}", params={'factory_id': 'huaxing'})
        assert current.status_code == 200, current.text
        current_rows = current.json()['parse_summary']['rows']
        assert len(current_rows) == 2 and all(item['schedule_identity_duplicate'] and item['can_operate_source'] for item in current_rows)
        assert business_snapshot() == before
        identities = []
        for value in current_rows:
            response = client.post(BASE + '/schedule-order-marks', json=dict(factory_id='huaxing',
                batch_id=duplicate['id'], source_sheet=value['source_sheet'], source_row=value['source_row'],
                marked=True, reason='逐行核实同SO的独立分批采购记录'))
            assert response.status_code == 200, response.text
            identities.append(response.json()['identity'])
        assert len(set(identities)) == 2
        after = business_snapshot()
        assert {key: value for key, value in after.items() if key != 'marks'} == {
            key: value for key, value in before.items() if key != 'marks'}


def test_inactive_customers_are_read_only_and_source_omission_requires_a_current_review(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client, monkeypatch)
        baseline = upload(client, 'base.xlsx', [row('INACTIVE1'), row('STAYS2')], declared=scope('FULL', '2026-10-10'))
        upload(client, 'missing.xlsx', [row('STAYS2')], declared=scope('FULL', '2026-10-11'))
        case = missing(client, 'INACTIVE1')
        before = business_snapshot()
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonCustomer
        from sqlalchemy import select
        with SessionLocal() as db:
            customer = db.scalar(select(CartonCustomer).where(CartonCustomer.factory_id == 'huaxing', CartonCustomer.customer_code == 'DICKIE'))
            customer.status = 'INACTIVE'; db.commit()
        work = workspace(client)
        assert next(item for item in work['customers'] if item['customer_code'] == 'DICKIE')['status'] == 'INACTIVE'
        readonly = detail(client, case['id'])
        assert readonly['can_review'] is False
        review(client, readonly, 'SOURCE_OMISSION', expected=422)
        assert detail(client, case['id'])['revision'] == case['revision']
        assert business_snapshot() == before
        with SessionLocal() as db:
            customer = db.scalar(select(CartonCustomer).where(CartonCustomer.factory_id == 'huaxing', CartonCustomer.customer_code == 'DICKIE'))
            customer.status = 'ACTIVE'; db.commit()
        restored, _ = review(client, detail(client, case['id']), 'SOURCE_OMISSION')
        assert restored['can_operate_source']
        reopened, _ = review(client, restored, 'REOPEN')
        assert reopened['status'] == 'PENDING' and not reopened['can_operate_source']
        body = order_body(baseline, 'INACTIVE1')
        assert client.post(BASE + '/orders', json=body).status_code == 409
        restored, _ = review(client, reopened, 'SOURCE_OMISSION')
        assert restored['can_operate_source'] and business_snapshot() == before
        created = client.post(BASE + '/orders', json=body)
        assert created.status_code == 201, created.text
