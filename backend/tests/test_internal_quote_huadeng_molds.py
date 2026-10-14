from io import BytesIO
from decimal import Decimal

import pytest
from openpyxl import Workbook

from app.services.internal_quote_import import parse_internal_quote_workbook


HEADERS = ["Item No. 模号", "Item Description 项目内容", "Material 原料", "Cavities 件",
           "Up套", "图片", "Mould Size(L*W*H) 工模尺寸", "Mould Prices 模价(RMB)",
           "钢材/钢料硬度", "产能", "机型", "胶件重量（g)", "Remark备注"]


def source_bytes(rows=None, merged=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "01"
    sheet.append(HEADERS)
    for row in rows or [
        ["M01", "前壳", "ABS", 2, 2, "", "", 1000, "", 3000, "18A", 12, "转水口"],
        [None, "后壳", "POM", 6, 2, "", "", None, "", None, None, 7.5, "2.5*3"],
        ["M02", "按钮", "POM", 8, 8, "", "", 2000, "", 4000, "18A", 1.2, ""],
        [None, "按钮", "", 8, 8, "", "", None, "", None, None, 1, ""],
        [], ["合计"], [None, "1、此报价不含税"], [None, "2、付款方式：预付"],
    ]:
        sheet.append(row)
    if merged:
        sheet.merge_cells("A4:A5")
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def test_huadeng_import_groups_continuations_and_preserves_each_source_part():
    parsed = parse_internal_quote_workbook(source_bytes(merged=True), "mold", factory_id="huadeng")
    molds = parsed.payload_fragment["molds"]
    assert parsed.row_count == 2
    assert [len(mold["parts"]) for mold in molds] == [2, 2]
    assert all(mold["molding_profile"] == "huadeng_parts_v1" for mold in molds)
    assert sum(Decimal(mold["cost_rmb"]) for mold in molds) == 3000
    parts = [part for mold in molds for part in mold["parts"]]
    assert len({part["part_id"] for part in parts}) == 4
    assert [part["source_row"] for part in parts] == [2, 3, 4, 5]
    assert [part["material_type"] for part in parts] == ["ABS", "POM", "POM", ""]
    assert parts[1]["remark"] == "2.5*3"
    assert Decimal(parts[1]["unit_net_weight_g"]) == Decimal("7.5")
    assert parts[1]["output_count"] == "6"
    assert Decimal(parts[1]["quantity"]) == 2
    assert parts[1]["machine_code"] == ""
    assert Decimal(parts[1]["target_output"]) == 0
    assert parts[3]["remark"] == ""


@pytest.mark.parametrize("factory", ["", "huakang-b", "huakang-d"])
def test_other_factories_keep_existing_row_semantics(factory):
    molds = parse_internal_quote_workbook(source_bytes(), "mold", factory_id=factory).payload_fragment["molds"]
    assert len(molds) == 4
    assert all("molding_profile" not in mold and "parts" not in mold for mold in molds)


def test_continuation_after_blank_cannot_inherit_previous_mold():
    with pytest.raises(ValueError, match="所属模号"):
        parse_internal_quote_workbook(source_bytes([
            ["M01", "前壳", "ABS", 2, 2, "", "", 1000, "", 3000, "18A", 12],
            [], [None, "孤立胶件", "POM", 2, 2, "", "", None, "", None, None, 3],
        ]), "mold", factory_id="huadeng")


def test_engineering_export_keeps_part_material_remark_and_weights_unmerged():
    import json
    from types import SimpleNamespace
    from openpyxl import load_workbook
    from app.services.internal_quote_excel import build_internal_quote_engineering_workbook

    payload = parse_internal_quote_workbook(source_bytes(), "mold", factory_id="huadeng").payload_fragment
    section = SimpleNamespace(department="engineering", payload_json=json.dumps(payload))
    quote = SimpleNamespace(factory_id="huadeng", workshop_name="测试", customer="测试", quote_no="TEST",
                            product_name="测试", version_label="V1", header_revision=1)
    book = load_workbook(BytesIO(build_internal_quote_engineering_workbook(quote, [section])))
    sheet = book["排摸表"]
    assert [sheet.cell(row, 2).value for row in range(6, 10)] == ["前壳", "后壳", "按钮", "按钮"]
    assert [sheet.cell(row, 9).value for row in range(6, 10)] == [12, 7.5, 1.2, 1]
    assert [sheet.cell(row, 12).value for row in range(6, 10)] == ["ABS", "POM", "POM", None]
    assert sheet.cell(6, 20).value == "转水口"
    assert sheet.cell(7, 20).value == "2.5*3"
    assert not any(merged.min_col == 12 and merged.min_row >= 6 for merged in sheet.merged_cells.ranges)
    book.close()


def test_continuation_images_survive_but_footer_images_are_excluded():
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image
    from PIL import Image as PillowImage

    book = load_workbook(BytesIO(source_bytes()))
    for anchor in ("F2", "F3", "F7"):
        data = BytesIO()
        PillowImage.new("RGB", (4, 4), "blue").save(data, format="PNG")
        data.seek(0)
        image = Image(data)
        image.anchor = anchor
        book.active.add_image(image)
    output = BytesIO()
    book.save(output)
    parsed = parse_internal_quote_workbook(output.getvalue(), "mold", factory_id="huadeng")
    assert [image.source_row for image in parsed.embedded_images] == [2, 3]
    assert parsed.payload_fragment["molds"][0]["parts"][1]["image_reference"]


def test_missing_source_sets_do_not_default_to_one():
    parsed = parse_internal_quote_workbook(source_bytes([
        ["M01", "待补胶件", "", None, None, "", "", 1000],
    ]), "mold", factory_id="huadeng")
    part = parsed.payload_fragment["molds"][0]["parts"][0]
    assert part["material_type"] == ""
    assert Decimal(part["quantity"]) == 0
    assert part["output_count"] == ""


def engineering_group(remarks=("", "", "转水口", "")):
    return {"molds": [{"mold_id": "physical", "mold_no": "M01", "item": "整模",
                       "molding_profile": "huadeng_parts_v1", "machine_code": "18A",
                       "target_output": "100", "net_weight_g": "999", "quantity": "2", "cost_rmb": "1000",
                       "parts": [{"part_id": f"p{i}", "name": "同名", "material_type": "ABS",
                                  "unit_net_weight_g": str(weight), "output_count": str(i + 2),
                                  "quantity": "2", "remark": remark}
                                 for i, (weight, remark) in enumerate(zip((2, 3, 5, 7), remarks))]}]}


def projected(engineering=None, existing=None):
    from app.services.internal_quote_prefill import prefill_molding_from_engineering
    return prefill_molding_from_engineering(engineering or engineering_group(), existing or {"injection_loss_rate_percent": "0"}, factory_id="huadeng")


def price(payload, factory="huadeng"):
    from app.services.internal_quote_calculator import calculate_section
    return calculate_section("molding", payload, {
        "material_prices": {"ABS|G": "454", "POM|P": "908"},
        "machine_prices": [{"range": "18A", "machine": "test", "shift_price_hkd": "200"},
                           {"range": "24A", "machine": "other", "shift_price_hkd": "400"}],
    }, "test", context={"factory_id": factory})


def fill_grades(payload):
    for row in payload["injection_lines"]:
        row["grade"] = "G"
    return payload


@pytest.mark.parametrize("value,expected", [(None, 100), ("", 100), ("  ", 100), (0, 100), ("0.0", 100), (250, 250)])
def test_child_capacity_only_blank_and_zero_inherit(value, expected):
    source = engineering_group(("",))
    source["molds"][0]["parts"][0]["target_output"] = value
    payload = fill_grades(projected(source))
    assert Decimal(str(payload["injection_lines"][0]["target_output"])) == expected
    assert price(payload)["status"] == "valid"


@pytest.mark.parametrize("field,value", [("material", "ABS"), ("material_type", "POM"),
                                         ("machine_code", "24A"), ("target_output", 250),
                                         ("remark", "转水口"), ("part_id", "replacement"),
                                         ("source_row", 20), ("image_attachment_ids", ["changed-image"])])
@pytest.mark.parametrize("status", ["approved", "pending_review"])
def test_child_source_only_change_invalidates_reviewed_molding(monkeypatch, field, value, status):
    import importlib
    from test_internal_quote_api import make_client, login, create_payload

    with make_client(monkeypatch) as client:
        login(client, "huadeng-owner", "sales_customer_owner", "sales-business", "huadeng")
        created = client.post("/api/internal-quotes", json={
            **create_payload(suffix="CHILD-DEPENDENCY", participating_sections=["engineering", "sales", "assembly", "molding"]),
            "factory_id": "huadeng", "workshop_code": "huadeng-workshop", "workshop_name": "华登", "workflow_mode": "direct_output",
        })
        assert created.status_code == 201, created.text
        quote_id = created.json()["id"]
        base = f"/api/internal-quotes/{quote_id}"
        sections = {row["department"]: row for row in created.json()["sections"]}
        source = engineering_group(("",))
        payload = projected(source)
        payload["injection_lines"][0]["grade"] = "750SW"
        saved = client.put(base + "/sections/save-all", json={"sections": [
            {"section_code": code, "revision": sections[code]["revision"], "payload": data}
            for code, data in (("engineering", source), ("molding", payload))]})
        assert saved.status_code == 200, saved.text
        sections = {row["department"]: row for row in client.get(base).json()["sections"]}
        before = sections["molding"]
        assert before["calculation_status"] == "valid"
        db_module = importlib.import_module("app.db")
        models = importlib.import_module("app.models.internal_quote")
        with db_module.SessionLocal() as db:
            row = db.query(models.InternalQuoteSection).filter_by(quote_id=quote_id, department="molding").one()
            row.status = status
            db.commit()
        changed = sections["engineering"]["payload"]
        changed["molds"][0]["parts"][0][field] = value
        saved = client.put(base + "/sections/engineering", json={"revision": sections["engineering"]["revision"], "payload": changed})
        assert saved.status_code == 200, saved.text
        with db_module.SessionLocal() as db:
            row = db.query(models.InternalQuoteSection).filter_by(quote_id=quote_id, department="molding").one()
            assert row.revision == before["revision"] + 1
            assert row.status == "rejected"
            assert row.calculation_status == row.dependency_status == "stale"
        current = next(row for row in client.get(base).json()["sections"] if row["department"] == "molding")
        recalculated = client.put(base + "/sections/molding", json={"revision": current["revision"], "payload": current["payload"]})
        assert recalculated.status_code == 200, recalculated.text
        with db_module.SessionLocal() as db:
            row = db.query(models.InternalQuoteSection).filter_by(quote_id=quote_id, department="molding").one()
            assert row.dependency_status == "current"
            assert row.calculation_status in {"valid", "blocked"}


@pytest.mark.parametrize("value,accepted", [(-1, False), ("abc", False), ("NaN", False), ("Infinity", False),
                                           (None, True), ("", True), (0, True), (250, True)])
def test_child_capacity_api_rejects_invalid_values_without_rewriting_saved_source(monkeypatch, value, accepted):
    from test_internal_quote_api import make_client, login, create_payload
    with make_client(monkeypatch) as client:
        login(client, "huadeng-owner", "sales_customer_owner", "sales-business", "huadeng")
        created = client.post("/api/internal-quotes", json={
            **create_payload(suffix="CHILD-CAPACITY", participating_sections=["engineering", "sales", "assembly", "molding"]),
            "factory_id": "huadeng", "workshop_code": "huadeng-workshop", "workshop_name": "华登", "workflow_mode": "direct_output",
        })
        assert created.status_code == 201, created.text
        base = f"/api/internal-quotes/{created.json()['id']}"
        section = next(row for row in created.json()["sections"] if row["department"] == "engineering")
        source = engineering_group(("",))
        source["molds"][0]["parts"][0]["target_output"] = value
        saved = client.put(base + "/sections/engineering", json={"revision": section["revision"], "payload": source})
        assert saved.status_code == (200 if accepted else 400), saved.text
        if not accepted:
            assert "产能" in saved.json()["detail"]
            after = next(row for row in client.get(base).json()["sections"] if row["department"] == "engineering")
            assert after["revision"] == section["revision"]
            assert after["payload"] == section["payload"]


@pytest.mark.parametrize("value", [-1, "abc", "NaN", "Infinity"])
def test_invalid_child_capacity_cannot_inherit_or_retain_a_previous_valid_price(value):
    from app.services.internal_quote_calculator import CalculationInputError, calculate_section
    source = engineering_group(("",))
    existing = fill_grades(projected(source))
    source["molds"][0]["parts"][0]["target_output"] = value
    payload = projected(source, existing)
    assert payload["injection_lines"][0]["target_output"] == value
    with pytest.raises(CalculationInputError):
        price(payload)
    with pytest.raises(CalculationInputError, match="产能"):
        calculate_section("engineering", source, {"fx": {"rmb_hkd": "1", "rmb_usd": "7"}}, "test", context={"factory_id": "huadeng"})


@pytest.mark.parametrize("count", [2, 4])
def test_huadeng_unmerged_body_cells_keep_column_typography_and_positional_borders(count):
    from copy import copy
    import json
    from types import SimpleNamespace
    from openpyxl import load_workbook
    from app.services.internal_quote_excel import ENGINEERING_WORKBOOK_TEMPLATE_PATH, build_internal_quote_engineering_workbook

    section = SimpleNamespace(department="engineering", payload_json=json.dumps(engineering_group(("",) * count)))
    payload = engineering_group(("",) * count)
    payload["molds"][0]["parts"][-1].update(quantity=0, output_count=0)
    section.payload_json = json.dumps(payload)
    quote = SimpleNamespace(factory_id="huadeng", workshop_name="测试", customer="测试", quote_no="STYLE",
                            product_name="测试", version_label="V1", header_revision=1)
    template = load_workbook(ENGINEERING_WORKBOOK_TEMPLATE_PATH)
    output = load_workbook(BytesIO(build_internal_quote_engineering_workbook(quote, [section])))
    try:
        source, sheet = template["排摸表"], output["排摸表"]
        for offset in range(count):
            row = offset + 6
            template_row = (6 if offset == 0 else 7) if count == 2 else (12 if offset == 0 else 16 if offset == count - 1 else 13)
            assert sheet.row_dimensions[row].height == source.row_dimensions[template_row].height
            for column in (3, 4, 5, 8, 12, 13, 14, 15, 16, 17):
                assert copy(sheet.cell(row, column).font) == copy(source.cell(17, column).font)
                assert copy(sheet.cell(row, column).alignment) == copy(source.cell(17, column).alignment)
                assert sheet.cell(row, column).number_format == source.cell(17, column).number_format
                assert copy(sheet.cell(row, column).border) == copy(source.cell(template_row, column).border)
        assert {str(value) for value in sheet.merged_cells.ranges if value.min_row >= 6 and value.max_row < count + 6} == {f"A6:A{count + 5}"}
        assert sheet.cell(count + 5, 11).value == sheet.cell(count + 5, 13).value == 0
        assert Decimal(str(sheet.cell(count + 5, 14).value)) == 0
    finally:
        template.close()
        output.close()


@pytest.mark.parametrize("remarks,jobs", [(('', '', '', ''), 1),
                                         (('转水口', '转水口'), 2),
                                         (('', '转水口', '', '转水口'), 3)])
def test_labor_jobs_charge_once_without_ghost_base_and_weights_do_not_multiply(remarks, jobs):
    payload = fill_grades(projected(engineering_group(remarks)))
    calculation = price(payload)
    assert calculation["status"] == "valid"
    lines = calculation["line_breakdown"]
    assert len(lines) == len(remarks)
    assert len({line["labor_job_id"] for line in lines}) == jobs
    assert sum(Decimal(line["molding_amount_hkd"]) for line in lines) == jobs
    assert sum(Decimal(line["material_amount_hkd"]) for line in lines) == sum((2, 3, 5, 7)[:len(remarks)])
    assert all(Decimal(row["quantity"]) == 1 for row in payload["injection_lines"])


def test_part_identity_reorder_edit_delete_detach_preserves_manual_work():
    from copy import deepcopy
    source = engineering_group()
    payload = fill_grades(projected(source))
    payload["injection_lines"][1].update(remark="人工备注", machine_code="24A", machine_name="chosen")
    source["molds"][0]["parts"].reverse()
    source["molds"][0]["parts"][2]["name"] = "已改名"
    refreshed = projected(source, payload)
    retained = next(row for row in refreshed["injection_lines"] if row["engineering_part_id"] == "p1")
    assert retained["item"] == "已改名"
    assert (retained["grade"], retained["machine_code"], retained["remark"]) == ("G", "24A", "人工备注")
    source["molds"][0]["parts"] = [part for part in source["molds"][0]["parts"] if part["part_id"] != "p1"]
    orphaned = projected(source, refreshed)
    orphan = next(row for row in orphaned["injection_lines"] if row["engineering_part_id"] == "p1")
    assert orphan["engineering_source_missing"] is True
    assert price(orphaned)["status"] == "blocked"
    orphan["engineering_sync_disabled"] = True
    detached = deepcopy(orphan)
    detached.pop("engineering_source_missing")
    after = projected(source, orphaned)
    assert next(row for row in after["injection_lines"] if row["engineering_part_id"] == "p1") == detached
    assert price(after)["status"] == "valid"
    assert len(after["injection_lines"]) == 4


def test_factory_isolation_no_legacy_auto_migration_and_invalid_shared_inputs_block():
    from app.services.internal_quote_prefill import prefill_molding_from_engineering
    from app.services.internal_quote_calculator import CalculationInputError
    payload = fill_grades(projected())
    with pytest.raises(CalculationInputError, match="仅适用于华登"):
        price(payload, "huakang-b")
    source = engineering_group()
    del source["molds"][0]["molding_profile"]
    assert len(prefill_molding_from_engineering(source, {}, factory_id="huadeng")["injection_lines"]) == 1
    payload["injection_lines"][1]["sets"] = "9"
    assert price(payload)["status"] == "blocked"


def test_mixed_materials_each_part_has_own_material_cost_and_source_missing_stays_pending():
    source = engineering_group()
    source["molds"][0]["parts"][1]["material_type"] = "POM"
    payload = fill_grades(projected(source))
    payload["injection_lines"][1]["grade"] = "P"
    calculated = price(payload)
    assert calculated["status"] == "valid"
    assert sum(Decimal(line["material_amount_hkd"]) for line in calculated["line_breakdown"]) == 20
    source["molds"][0]["parts"][0]["material_type"] = ""
    missing = projected(source)
    assert missing["injection_lines"][0]["material"] == ""
    assert missing["injection_lines"][0]["grade"] == ""
    assert price(missing)["status"] == "blocked"


def test_export_formulas_charge_jobs_once_and_never_merge_across_transfer():
    from app.services.internal_quote_excel import _apply_molding_row_formulas, _merge_molding_labor_jobs
    lines = price(fill_grades(projected(engineering_group(("", "", "转水口", "")))))["line_breakdown"]
    sheet = Workbook().active
    for index, line in enumerate(lines, start=9):
        sheet.cell(index, 8).value = float(line["sets"])
        sheet.cell(index, 9).value = float(line["target_output"])
        _apply_molding_row_formulas(sheet, index, line, {}, {})
        sheet.cell(index, 12).value = f"=J{index}*1.15"
    assert sheet["J10"].value == "=0"
    assert sheet["J12"].value == "=0"
    assert sheet["J9"].value != "=0" and sheet["J11"].value != "=0"
    _merge_molding_labor_jobs(sheet, lines, 9)
    assert {str(value) for value in sheet.merged_cells.ranges} == {"J9:J10", "L9:L10"}


def test_huadeng_api_save_preview_copy_and_issued_history_keep_exact_part_costs(monkeypatch):
    import importlib
    from copy import deepcopy
    from openpyxl import load_workbook
    from test_internal_quote_api import make_client, login, create_payload
    from test_internal_quote_alternatives import fill, fork

    with make_client(monkeypatch) as client:
        login(client, "huadeng-owner", "sales_customer_owner", "sales-business", "huadeng")
        created = client.post("/api/internal-quotes", json={
            **create_payload(suffix="HUADENG-PARTS", participating_sections=["engineering", "sales", "assembly", "molding"]),
            "factory_id": "huadeng", "workshop_code": "huadeng-workshop", "workshop_name": "华登", "workflow_mode": "direct_output",
        })
        assert created.status_code == 201, created.text
        base = f"/api/internal-quotes/{created.json()['id']}"
        ready = fill(client, created.json())
        sections = {row["department"]: row for row in ready["sections"]}
        source = engineering_group()
        source.update(mold_allocation_enabled=False, amortization_qty=5000)
        payload = projected(source)
        for row in payload["injection_lines"]:
            row["grade"] = "750SW"
        drafts = [{"section_code": code, "revision": sections[code]["revision"], "payload": value}
                  for code, value in (("engineering", source), ("molding", payload))]
        saved = client.put(base + "/sections/save-all", json={"sections": drafts})
        assert saved.status_code == 200, saved.text
        current = client.get(base).json()
        sections = {row["department"]: row for row in current["sections"]}
        molding_section = sections["molding"]
        assert molding_section["calculation_status"] == "valid"
        calculation = molding_section["calculation"]
        assert len(calculation["line_breakdown"]) == 4
        assert len({row["labor_job_id"] for row in calculation["line_breakdown"]}) == 2
        preview = client.post(base + "/cost-preview", json={"drafts": [{"section_code": "molding", "revision": molding_section["revision"], "payload": molding_section["payload"]}]})
        assert preview.status_code == 200, preview.text
        stale = client.put(base + "/sections/molding", json={"revision": molding_section["revision"] - 1, "payload": molding_section["payload"]})
        assert stale.status_code == 409

        # Read projection is advisory and never rewrites persisted draft evidence.
        db_module = importlib.import_module("app.db")
        models = importlib.import_module("app.models.internal_quote")
        def persisted():
            with db_module.SessionLocal() as db:
                row = db.query(models.InternalQuoteSection).filter_by(quote_id=created.json()["id"], department="molding").one()
                return row.revision, row.payload_json, row.calculation_json
        before_read = persisted()
        client.get(base)
        client.get(base + "/summary")
        assert persisted() == before_read
        issued = client.post(base + "/direct-issue", json={"revision": current["header_revision"]})
        assert issued.status_code == 200, issued.text
        download = base + f"/exports/{issued.json()['id']}/download"
        frozen_bytes = client.get(download).content
        book = load_workbook(BytesIO(frozen_bytes))
        assert any("J9:J10" in {str(value) for value in sheet.merged_cells.ranges} for sheet in book.worksheets)
        book.close()
        frozen = client.get(base).json()
        frozen_molding = next(row for row in frozen["sections"] if row["department"] == "molding")
        assert frozen_molding["payload"] == molding_section["payload"]
        copied = fork(client, frozen, "version")
        copied_molding = next(row for row in copied["sections"] if row["department"] == "molding")
        assert copied_molding["payload"] == frozen_molding["payload"]
        assert copied_molding["calculation"]["totals"] == calculation["totals"]
        assert client.get(download).content == frozen_bytes
        denied = client.put(base + "/sections/engineering", json={"revision": sections["engineering"]["revision"], "payload": deepcopy(source)})
        assert denied.status_code == 409


def test_huadeng_api_import_profile_is_bound_to_persisted_factory(monkeypatch):
    from test_internal_quote_api import make_client, login, create_payload
    with make_client(monkeypatch) as client:
        for factory in ("huadeng", "huaxing"):
            login(client, "profile-" + factory, "sales_customer_owner", "sales-business", factory)
            created = client.post("/api/internal-quotes", json={**create_payload(suffix="PROFILE-" + factory, participating_sections=["engineering", "sales", "assembly", "molding"]),
                                  "factory_id": factory, "workshop_code": factory + "-workshop", "workflow_mode": "direct_output"})
            assert created.status_code == 201, created.text
            base = f"/api/internal-quotes/{created.json()['id']}"
            preview = client.post(base + "/imports/mold/preview", files={"file": ("arbitrary.xlsx", source_bytes())})
            assert preview.status_code == 201, preview.text
            confirmed = client.post(base + f"/imports/{preview.json()['batch_id']}/confirm", json={"revision": preview.json()["target_revision"], "mode": "replace"})
            assert confirmed.status_code == 200, confirmed.text
            section = confirmed.json()["section"]
            molds = section["payload"]["molds"]
            assert len(molds) == (2 if factory == "huadeng" else 4)
            assert all((row.get("molding_profile") == "huadeng_parts_v1") == (factory == "huadeng") for row in molds)
            if factory == "huadeng":
                original_ids = [(mold["mold_id"], [part["part_id"] for part in mold["parts"]]) for mold in molds]
                current_molding = next(row for row in client.get(base).json()["sections"] if row["department"] == "molding")
                for index, row in enumerate(current_molding["payload"]["injection_lines"]):
                    row.update(grade=f"manual-{index}", machine_code="24A", remark=f"own-{index}")
                saved_molding = client.put(base + "/sections/molding", json={"revision": current_molding["revision"], "payload": current_molding["payload"]})
                assert saved_molding.status_code == 200, saved_molding.text
                manual_rows = saved_molding.json()["payload"]["injection_lines"]
                repeated = client.post(base + "/imports/mold/preview", files={"file": ("same-content-renamed.xlsx", source_bytes())})
                assert repeated.status_code == 201, repeated.text
                repeated_confirm = client.post(base + f"/imports/{repeated.json()['batch_id']}/confirm", json={"revision": repeated.json()["target_revision"], "mode": "replace"})
                assert repeated_confirm.status_code == 200, repeated_confirm.text
                reimported = repeated_confirm.json()["section"]["payload"]["molds"]
                assert [(mold["mold_id"], [part["part_id"] for part in mold["parts"]]) for mold in reimported] == original_ids
                after = next(row for row in client.get(base).json()["sections"] if row["department"] == "molding")["payload"]["injection_lines"]
                assert len(after) == 4
                assert all(not row.get("engineering_source_missing") for row in after)
                assert [(row["line_id"], row["grade"], row["machine_code"], row["remark"]) for row in after] == [
                    (row["line_id"], row["grade"], row["machine_code"], row["remark"]) for row in manual_rows]
            if factory != "huadeng":
                spoofed = {**engineering_group(), "factory_id": "huadeng"}
                rejected = client.put(base + "/sections/engineering", json={"revision": section["revision"], "payload": spoofed})
                assert rejected.status_code == 400


def test_reimport_identity_never_guesses_moved_or_ambiguous_source_rows():
    from copy import deepcopy
    from app.services.internal_quote_prefill import reuse_huadeng_import_identities
    original = parse_internal_quote_workbook(source_bytes(), "mold", factory_id="huadeng").payload_fragment
    incoming = parse_internal_quote_workbook(source_bytes(), "mold", factory_id="huadeng").payload_fragment
    incoming["molds"][0]["parts"][1]["source_row"] = 99
    incoming_part_id = incoming["molds"][0]["parts"][1]["part_id"]
    result = reuse_huadeng_import_identities(incoming, original, factory_id="huadeng")
    assert result["molds"][0]["mold_id"] == original["molds"][0]["mold_id"]
    assert result["molds"][0]["parts"][1]["part_id"] == incoming_part_id
    ambiguous = deepcopy(original)
    ambiguous["molds"].append(deepcopy(original["molds"][0]))
    result = reuse_huadeng_import_identities(incoming, ambiguous, factory_id="huadeng")
    assert result["molds"][0]["mold_id"] == incoming["molds"][0]["mold_id"]
