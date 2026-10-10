"""Shared Shixin input parsing must not share factory histories or approvals."""
from copy import deepcopy
from io import BytesIO

import openpyxl
import pytest
from fastapi import HTTPException
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from app.api import customer_order as api
from app.services import customer_order_unified as service
from test_customer_order_huaxing_unified import template, data
from test_customer_order_ledger import engine, db, client, user
from app.services import customer_order_ledger as ledger
from app.services.auth import get_current_user


def seasons_po():
    text = '''Customer Name: SEASONS USA
Order Date: 2026/10/01
MPO1609793
W85340 OC: ZE618957 SHIPDATE: 2026/11/10 120 2.50 0 300.00
TEST PUMPKIN
Packing: Outer: 12 PCS'''
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                             NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    stream.set_data(('BT /F1 9 Tf 12 TL 30 760 Td ' + ' T* '.join(f'({line}) Tj' for line in text.splitlines()) + ' ET').encode())
    page[NameObject('/Contents')] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return [('MPO1609793.pdf', output.getvalue())]


def arguments(factory):
    wb = template('seasons')
    if factory == 'huaxing':
        # Existing Huaxing SEASONS keeps its older aligned-boundary writer.
        for name in ['接单表', '正单评审表']:
            for row in range(4, wb[name].max_row + 1):
                wb[name].cell(row, 1).value = None
            wb[name]['A8'] = '取消单'
    ws = wb['ITEM表']
    ws['E4'], ws['H4'], ws['I4'] = factory + '-OLD', 'W85340', factory + '-产品'
    ws['K4'], ws['L4'] = 24, 12
    return dict(customer_code='seasons', factory_id=factory, received_date='2026-10-06',
                po_files=seasons_po(), schedule_file_name=factory + '.xlsx', schedule_content=data(wb))


@pytest.mark.parametrize('factory', ['huaxing', 'huakang-d'])
def test_seasons_preview_export_use_only_current_factory_workbook(factory):
    api._ensure_customer_factory('seasons', factory)
    assert '.pdf' in api._get_mapped_customer_spec('seasons', factory).po_extensions
    args = arguments(factory)
    preview = service.create_unified_customer_preview(**args)
    assert preview['factory_id'] == factory
    row = preview['rows'][0]
    assert (row['reference_no'], row['po_no'], row['product_no'], row['quantity']) == (
        'ZE618957', 'MPO1609793', 'W85340', '120')
    assert row['product_name_zh'] == 'TEST PUMPKIN'
    keys = {i['skip_key'] for r in preview['rows'] for i in r['issues'] if i.get('skip_key')}
    output, _, exported = service.export_unified_customer_schedule(**args, skipped_issue_keys=keys)
    assert exported['factory_id'] == factory
    wb = openpyxl.load_workbook(BytesIO(output), data_only=False)
    assert wb['ITEM表']['E4'].value == factory + '-OLD'
    assert wb['ITEM表']['E5'].value == 'ZE618957'
    assert wb['ITEM表']['F5'].value == 'MPO1609793'
    assert wb['ITEM表']['I4'].value == factory + '-产品'
    assert wb['ITEM表']['I5'].value == 'TEST PUMPKIN'
    assert wb['ITEM表']['K5'].value == 120
    assert wb['ITEM表']['K9'].value is None
    if factory == 'huakang-d':
        assert preview['target_template'] == row['target_template'] == 'HEYUAN_BUSINESS_UNIFIED_REGIONAL_V3'
        for name in ['接单表', '正单评审表']:
            assert wb[name]['D4'].value == 'ZE618957'
            assert wb[name]['E4'].value == 'MPO1609793'
            assert "'ITEM表'!K5" in wb[name]['J4'].value
    assert openpyxl.load_workbook(BytesIO(args['schedule_content']))['ITEM表']['E5'].value is None


@pytest.mark.parametrize('factory', ['huakang-a', 'huakang-b', 'huakang-c', 'huadeng'])
def test_seasons_other_factories_still_rejected(factory):
    with pytest.raises(HTTPException) as error:
        api._ensure_customer_factory('seasons', factory)
    assert error.value.status_code == 400
    with pytest.raises(service.CustomerOrderUnifiedError, match='不属于当前厂区'):
        service.create_unified_customer_preview(**arguments(factory))


def test_seasons_approval_cannot_be_reused_across_factories():
    args = arguments('huaxing')
    preview = service.create_unified_customer_preview(**args)
    fingerprint = api._preview_fingerprint(preview, args['received_date'])
    other = deepcopy(preview)
    other['factory_id'] = 'huakang-d'
    with pytest.raises(HTTPException) as error:
        api._ensure_preview_fingerprint(other, args['received_date'], fingerprint)
    assert error.value.status_code == 409


def test_seasons_factory_history_and_quantity_conflict_remain_isolated():
    args = arguments('huaxing')
    wb = openpyxl.load_workbook(BytesIO(args['schedule_content']))
    wb['ITEM表']['D4'], wb['ITEM表']['E4'] = 'ZE618957', 'ZE618957'
    wb['ITEM表']['F4'] = 'MPO1609793'
    # Different quantity is a modification blocker, not a confirmable duplicate.
    args['schedule_content'] = data(wb)
    hx = service.create_unified_customer_preview(**args)
    hd = service.create_unified_customer_preview(**arguments('huakang-d'))
    conflict = next(i for i in hx['rows'][0]['issues'] if i['code'] == 'existing_quantity_conflict')
    assert not conflict.get('can_skip')
    assert not any(i['code'] in {'existing_quantity_conflict', 'duplicate_existing_order'} for i in hd['rows'][0]['issues'])


def test_seasons_factory_options_include_each_factory_once():
    assert api.CUSTOMER_FACTORY_OPTIONS['seasons'] == ('huaxing', 'huakang-d')
    assert api.CUSTOMER_FACTORY_OPTIONS['disney'] == ('huaxing', 'huakang-d')


def test_seasons_persisted_orders_and_read_permissions_are_factory_scoped(client, db):
    ids = {}
    for factory in ['huaxing', 'huakang-d']:
        payload = service.create_unified_customer_preview(**arguments(factory))
        imported = ledger.import_preview(db, preview=payload, files=[
            ('MPO1609793.pdf', seasons_po()[0][1], 'formal')],
            source_kind='formal', actor='tester', reason='已核对订单', controls={})
        db.commit()
        ids[factory] = imported['items'][0]['id']
    assert ids['huaxing'] != ids['huakang-d']
    with pytest.raises(HTTPException) as error:
        ledger.get_line(db, 'huakang-d', ids['huaxing'])
    assert error.value.status_code == 404
    for factory in ['huaxing', 'huakang-d']:
        client.test_app.dependency_overrides[get_current_user] = lambda factory=factory: user(['read'], factory=factory)
        response = client.get('/api/customer-order-ledger/lines', params={'factory_id': factory})
        assert response.status_code == 200
        assert [r['id'] for r in response.json()['items']] == [ids[factory]]
        customers = client.get('/api/customer-order-ledger/history/customers', params={'factory_id': factory})
        assert len([r for r in customers.json()['items'] if r['code'] == 'seasons']) == 1
        other = 'huaxing' if factory == 'huakang-d' else 'huakang-d'
        assert client.get('/api/customer-order-ledger/lines', params={'factory_id': other}).status_code == 403
