"""Cutover acceptance against the published workbooks and isolated API database."""
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from threading import Barrier

from openpyxl import load_workbook

from test_molding_sample_api import make_client,login_as
from test_carton_history_identity_api import _customer
from test_carton_history_preview_api import upload as order_upload
from test_carton_opening_preview_api import upload as opening_upload,setup as _base_opening_setup,BATCH

BASE="/api/carton-procurement"
ROOT=Path(__file__).resolve().parents[2]


def opening_setup(client, monkeypatch):
    from datetime import datetime
    from zoneinfo import ZoneInfo
    from app.services import carton_procurement
    result = _base_opening_setup(client, monkeypatch)
    # The published template now posts its August document dates, not BATCH's date.
    monkeypatch.setattr(carton_procurement, "business_now", lambda: datetime(2026,9,1,10,tzinfo=ZoneInfo("Asia/Shanghai")))
    return result


def examples_as_input(path,main_sheet,example_rows):
    workbook=load_workbook(ROOT/path)
    for target_row,source_row in enumerate(example_rows,3):
        for col,cell in enumerate(workbook["填写示例"][source_row],1):
            workbook[main_sheet].cell(target_row,col,cell.value)
    stream=BytesIO();workbook.save(stream);return stream.getvalue()


def orders(client):
    response=client.get(f"{BASE}/orders",params={"factory_id":"huaxing"})
    assert response.status_code==200,response.text
    return response.json()["items"]


def test_actual_eight_field_history_template_import_edit_lock_issue_preserves_unknowns_and_no_false_master():
    from pytest import MonkeyPatch
    with MonkeyPatch.context() as monkeypatch, make_client(monkeypatch) as client:
        login_as(client,"admin");_customer(client,"迪奇")
        content=examples_as_input("public/templates/carton-history-order-import-template.xlsx","历史订单导入",[6,7])
        response=order_upload(client,content);assert response.status_code==200,response.text
        preview=response.json();assert preview["errors"]==[]
        assert preview["group_count"]==2 and preview["line_count"]==4
        assert preview["incomplete_count"]==2
        assert {row["item_no"] for row in preview["orders"]}=={"000203307004"}
        result=order_upload(client,content,preview=False,fingerprint=preview["source_fingerprint"])
        assert result.status_code==201,result.text
        saved=orders(client)
        first=next(row for row in saved if row["contract_no"]=="0010001234/200")
        other=next(row for row in saved if row["contract_no"]=="0010001234/300")
        assert first["product_order_quantity"] is None
        assert [(row["packaging_type"],Decimal(row["required_quantity"])) for row in other["lines"]]==[("内箱",200),("外箱",100),("卡纸",100)]
        assert all(row["usage_quantity"] is None for order in saved for row in order["lines"])
        assert first["status"] == "PENDING_SUPPLIER"
        context=client.get(f"{BASE}/orders/{first['order_no']}/purchase-order-context",params={"factory_id":"huaxing"}).json()
        assert context["historical_baseline"] and context["pending_type"] == "NONE"
        assert context["issues"] == [] and not context["can_generate"]
        master=client.get(f"{BASE}/master-data",params={"factory_id":"huaxing"})
        assert master.status_code==200,master.text
        assert not [record for record in master.json()["records"] if record["kind"]=="CONFIG"]
        assert client.get(f"{BASE}/inventory/movements",params={"factory_id":"huaxing"}).json()["total"]==0
        assert all(Decimal(line["received_quantity"])==0 for order in orders(client) for line in order["lines"])


def test_actual_opening_template_concurrent_confirmation_one_post_and_row_audit_preserved(monkeypatch):
    with make_client(monkeypatch) as client:
        opening_setup(client,monkeypatch)
        loc=client.post(f"{BASE}/inventory/locations",json={"factory_id":"huaxing","warehouse":"A","bin_code":"A2"})
        assert loc.status_code==201,loc.text
        content=examples_as_input("public/templates/carton-history-inventory-import-template.xlsx","期初库存导入",[6,7,8,9])
        response=opening_upload(client,content);assert response.status_code==200,response.text
        preview=response.json();assert preview["errors"]==[]
        assert preview["skipped_count"]==1 and preview["missing_price_count"]==1
        assert {row["item_no"] for row in preview["rows"]}=={"000203307004","000203307005"}
        barrier=Barrier(2)
        def confirm():
            barrier.wait(timeout=10)
            return opening_upload(client,content,preview=False,token=preview["source_fingerprint"])
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses=[future.result(timeout=30) for future in [pool.submit(confirm),pool.submit(confirm)]]
        assert [response.status_code for response in responses]==[201,201],[response.text for response in responses]
        assert sorted(response.json()["imported_count"] for response in responses)==[0,3]
        assert sorted(response.json()["duplicate"] for response in responses)==[False,True]
        physical=client.get(f"{BASE}/inventory/positions",params={"factory_id":"huaxing"}).json()
        assert len(physical)==3
        assert sum(Decimal(row["balance"]) for row in physical if row["unit"]=="个")==300
        assert sum(Decimal(row["balance"]) for row in physical if row["unit"]=="张")==50
        assert next(row for row in physical if row["packaging_type"]=="内箱")["cost_amount"] is None
        assert Decimal(next(row for row in physical if row["packaging_type"]=="外箱")["cost_amount"])==250
        assert Decimal(next(row for row in physical if row["packaging_type"]=="平卡")["cost_amount"])==15
        from app import db as app_db
        from app.models.carton_procurement import CartonAuditEvent
        from sqlalchemy import select
        import json
        with app_db.SessionLocal() as db:
            audits=list(db.scalars(select(CartonAuditEvent).where(CartonAuditEvent.event_type=="HISTORY_INVENTORY_IMPORTED")))
            assert len(audits)==1
            detail=json.loads(audits[0].detail_json)
            assert len(detail["rows"])==3
            assert detail["rows"][0]["legacy_row_no"]=="OPEN-001"
            assert detail["rows"][0]["document_no"]=="OPEN-EXAMPLE"
            assert detail["rows"][0]["source_row"]==3
            assert detail["rows"][0]["movement_id"]


def test_opening_and_demand_previews_are_factory_scoped_and_have_no_partial_write_on_customer_error(monkeypatch):
    with make_client(monkeypatch) as client:
        opening_setup(client,monkeypatch)
        content=examples_as_input("public/templates/carton-history-order-import-template.xlsx","历史订单导入",[6])
        preview=order_upload(client,content).json()
        foreign=order_upload(client,content,preview=False,fingerprint=preview["source_fingerprint"],factory="huadeng")
        assert foreign.status_code in (409,422),foreign.text
        assert orders(client)==[]
        _customer(client,"迪奇",factory="huadeng")
        foreign_preview=order_upload(client,content,factory="huadeng").json()
        assert foreign_preview["source_fingerprint"]!=preview["source_fingerprint"]
        foreign=order_upload(client,content,preview=False,fingerprint=preview["source_fingerprint"],factory="huadeng")
        assert foreign.status_code==409,foreign.text
        assert client.get(f"{BASE}/orders",params={"factory_id":"huadeng"}).json()["total"]==0
        mixed=load_workbook(BytesIO(examples_as_input("public/templates/carton-history-order-import-template.xlsx","历史订单导入",[6,7])))
        mixed["历史订单导入"].cell(4,1,"尚未建立的客户")
        stream=BytesIO();mixed.save(stream)
        bad_preview=order_upload(client,stream.getvalue()).json()
        assert bad_preview["errors"] and len(bad_preview["orders"])==1
        rejected=order_upload(client,stream.getvalue(),preview=False,fingerprint=bad_preview["source_fingerprint"])
        assert rejected.status_code==422,rejected.text
        assert orders(client)==[]
        opening_content=examples_as_input("public/templates/carton-history-inventory-import-template.xlsx","期初库存导入",[6])
        opening_preview=opening_upload(client,opening_content).json()
        assert opening_preview["errors"]==[]
        foreign_bin=client.post(f"{BASE}/inventory/locations",json={"factory_id":"huadeng","warehouse":"A","bin_code":"A1"})
        assert foreign_bin.status_code==201,foreign_bin.text
        import json
        rejected=client.post(f"{BASE}/inventory/history-imports",params={"factory_id":"huadeng"},
            files={"file":("opening.xlsx",opening_content)},data={"options":json.dumps(BATCH),"expected_preview_fingerprint":opening_preview["source_fingerprint"]})
        assert rejected.status_code==409,rejected.text
        for factory in ("huaxing","huadeng"):
            assert client.get(f"{BASE}/inventory/movements",params={"factory_id":factory}).json()["total"]==0


def test_legacy_unitless_opening_cannot_be_reimported_as_guessed_cm_stock(monkeypatch):
    from test_carton_procurement_api import _history_inventory_workbook_bytes
    from test_carton_opening_preview import workbook
    with make_client(monkeypatch) as client:
        opening_setup(client,monkeypatch)
        response=client.post(f"{BASE}/inventory/locations",json={"factory_id":"huaxing","warehouse":"默认仓","bin_code":"OLD1"})
        assert response.status_code==201,response.text
        row=["PRE-CUTOVER","迪奇","","PO-LEGACY","ITEM-LEGACY","外箱","A33","30*20*15","个",10,2,"CNY","OLD1","2026-08-01","LEGACY-DOC","",None]
        result=opening_upload(client,_history_inventory_workbook_bytes([row]),preview=False,options=None)
        assert result.status_code==201,result.text
        content=workbook([["PO-LEGACY","ITEM-LEGACY","A33外箱",30,20,15,"OLD1",10,2]])
        preview=opening_upload(client,content,options={**BATCH,"warehouse":"默认仓"}).json()
        assert any("未明确尺寸单位" in error for error in preview["errors"])
        rejected=opening_upload(client,content,preview=False,options={**BATCH,"warehouse":"默认仓"},token=preview["source_fingerprint"])
        assert rejected.status_code==422,rejected.text
        assert client.get(f"{BASE}/inventory/movements",params={"factory_id":"huaxing"}).json()["total"]==1
