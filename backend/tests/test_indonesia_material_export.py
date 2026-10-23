"""Regression examples for the confirmed one-material-row packing output."""
import base64
from copy import copy, deepcopy
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile

import openpyxl
import pytest
from fastapi import HTTPException
from lxml import etree as ET
from PIL import Image

from app.services.indonesia_material_calculations import packaging_totals
from app.services.indonesia_material_excel import (export_shipment, export_shipment_with_summary, NS,
    SHEET_NAME, TEMPLATE_SHEET_NAME, TEMPLATE, SHEET_PART, DRAWING_PART, CALC_NAME, CALC_PART,
    prepare_template, relationship_path, target_path)


def source(code, *, price='.1', supplier='供应商甲', currency='USD', net='1', item='ITEM-A', contract='CONTRACT-A', po='PO-A'):
    photo = BytesIO(); Image.new('RGB', (4, 4), 'green').save(photo, 'PNG')
    material = dict(code=code, name='测试物料 '+code, english_name='Test material '+code,
                    specification='TEST', category='辅助材料', unit='PCS', customs_unit='KGM', net_weight_kg=net,
                    billing_weight_kg='.002', usage='装配', department='装配部', export_company='测试出口公司',
                    image_data='data:image/png;base64,'+base64.b64encode(photo.getvalue()).decode())
    supply = dict(material_id=code, supplier=supplier, currency=currency, price=price, billing_basis='quantity',
                  billing_factor='1', purchase_date='2026-10-01', purchase_no=po)
    return dict(material=material, supply=supply, item_no=item, contract_no=contract)


def package(label, items, *, cartons=1, weight='1', dims=('27', '14.5', '27.3')):
    return dict(label=label, cartons=cartons, carton_weight_kg=weight, length_cm=dims[0], width_cm=dims[1], height_cm=dims[2],
                pallet='不计入重量的卡板', items=[dict(allocation_id=key, normal_quantity=str(normal), spare_quantity=str(spare))
                                             for key, normal, spare in items])


def frozen(sources, packages):
    lines = deepcopy(sources)
    for key, line in lines.items():
        line.update(allocation_id=key, normal_quantity='0', spare_quantity='0')
    for group in packages:
        for item in group['items']:
            line = lines[item['allocation_id']]
            for field in ('normal_quantity', 'spare_quantity'):
                line[field] = str(Decimal(line[field])+Decimal(item[field]))
    return dict(recipient='RRI', customer='测试客户', remarks='', snapshot_lines=list(lines.values()),
                packing_totals=packaging_totals(packages, lines))


def detail_book(source, **options):
    """Read only the tested views; preservation of other tabs is checked as OOXML.

    The unchanged source invoice contains 246,493 styled cells. Excluding it
    from this disposable reader avoids repeatedly loading unrelated content.
    The actual exported bytes are never modified by this test helper.
    """
    if not hasattr(source, 'read'):
        source = BytesIO(source.read_bytes())
    with ZipFile(source) as archive:
        parts = {p: archive.read(p) for p in archive.namelist()}
    book = ET.fromstring(parts['xl/workbook.xml'])
    sheets = book.find('s:sheets', NS)
    old_index = str(next(i for i, tab in enumerate(sheets) if tab.get('name') == SHEET_NAME))
    sheets[:] = [tab for tab in sheets if tab.get('name') in {SHEET_NAME, CALC_NAME}]
    names = book.find('s:definedNames', NS)
    if names is not None:
        names[:] = [n for n in names if n.get('localSheetId') == old_index]
        for name in names:
            name.set('localSheetId', '0')
    view = book.find('s:bookViews/s:workbookView', NS)
    if view is not None:
        view.set('activeTab', '0'); view.set('firstSheet', '0')
    parts['xl/workbook.xml'] = ET.tostring(book)
    out = BytesIO()
    with ZipFile(out, 'w') as archive:
        for name, value in parts.items():
            archive.writestr(name, value)
    out.seek(0)
    return openpyxl.load_workbook(out, **options)


def output(data):
    content = export_shipment(data)
    cached = detail_book(BytesIO(content), data_only=True)[SHEET_NAME]
    formulas = detail_book(BytesIO(content), data_only=False)[SHEET_NAME]
    return content, cached, formulas


def test_frozen_supplier_display_matches_excel_rounding_and_distinct_material_counts():
    data = frozen({'a': source('A', price='.035'), 'a2': source('A', price='0'), 'b': source('B', price='.005'),
                   'c': source('C', price='2', supplier='供应商乙', currency='CNY')}, [
        package('AB', [('a', 1, 10), ('a2', 1, 0), ('b', 1, 0)]), package('C', [('c', 1, 0)])])
    data['snapshot_lines'][0]['supply']['billing_factor'] = '7'
    # Old source amount can fall below .005 due to division before multiplication.
    data['snapshot_lines'][0]['amount'] = '0.0049999999999999999999999999999999999999'
    before = deepcopy(data)
    content, summary = export_shipment_with_summary(data)
    cached = detail_book(BytesIO(content), data_only=True)[SHEET_NAME]
    assert cached['AP4'].value == .01
    assert summary == [dict(name='供应商甲', currency='USD', materials=2, amount='0.02'),
                       dict(name='供应商乙', currency='CNY', materials=1, amount='2.00')]
    assert data == before


def test_full_and_tail_cartons_are_one_material_row_with_actual_loadings():
    data = frozen({'a': source('A', net='.001')}, [
        package('整箱', [('a', 109200, 0)], cartons=52, weight='.1'),
        package('尾箱', [('a', 1800, 0)], weight='.1')])
    content, cached, formulas = output(data)
    assert cached['AZ4'].value == 'A' and cached['AZ5'].value is None
    assert cached['AU4'].value == '52箱×2100＋1箱×1800'
    assert cached['L4'].value == 111000 and cached['AT4'].value == 53
    assert cached['Q4'].value == 111 and cached['P4'].value == pytest.approx(116.3)
    assert formulas['P4'].value.startswith("=Q4+SUM('出货计算'!M")
    assert formulas['BB4'].value.startswith("='出货计算'!Y")
    assert cached['BB4'].value == '52箱×2.200000 kg＋1箱×1.900000 kg'
    assert cached['S4'].value == float(data['packing_totals']['cbm'])
    with ZipFile(BytesIO(content)) as archive:
        drawing = ET.fromstring(archive.read(DRAWING_PART))
        assert len(drawing.findall('.//xdr:pic', NS)) == 1


def test_mixed_carton_merges_count_and_volume_but_preserves_each_material_gross_weight():
    data = frozen({'a': source('A'), 'b': source('B')}, [package('混箱1', [('a', 6, 0), ('b', 3, 1)])])
    _, cached, formulas = output(data)
    assert {str(r) for r in formulas.merged_cells.ranges} >= {'AT4:AT5', 'S4:S5'}
    assert cached['AT4'].value == 1 and cached['AT5'].value is None
    assert cached['S4'].value == float(data['packing_totals']['cbm']) and cached['S5'].value is None
    assert cached['P4'].value == pytest.approx(6.6) and cached['P5'].value == pytest.approx(4.4)
    assert cached['BB4'].value == pytest.approx(6.6) and cached['BB5'].value == pytest.approx(4.4)
    assert cached['AP4'].value == .6 and cached['AP5'].value == .3
    assert cached['AQ4'].value == .9 and cached['AQ5'].value is None
    assert '备品 1 PCS，免费' in cached['BE5'].value


def test_overlapping_mixed_groups_merge_once_including_dedicated_cartons():
    data = frozen({'a': source('A'), 'd': source('D'), 'b': source('B'), 'c': source('C')}, [
        package('AB混箱', [('a', 6, 0), ('b', 4, 0)]),
        package('BC混箱', [('b', 6, 0), ('c', 4, 0)]),
        package('A独装', [('a', 3, 0)]), package('D独装', [('d', 2, 0)])])
    content, cached, formulas = output(data)
    assert [cached[f'AZ{r}'].value for r in range(4, 8)] == ['A', 'B', 'C', 'D']
    assert {str(r) for r in formulas.merged_cells.ranges} >= {'AT4:AT6', 'S4:S6'}
    assert cached['AT4'].value == 3 and cached['AT7'].value == 1
    assert all(cached[f'AT{r}'].value is None for r in (5, 6))
    assert sum(cached[f'AT{r}'].value or 0 for r in range(4, 8)) == data['packing_totals']['cartons']
    assert sum(cached[f'S{r}'].value or 0 for r in range(4, 8)) == pytest.approx(float(data['packing_totals']['cbm']))
    assert sum(cached[f'P{r}'].value for r in range(4, 8)) == pytest.approx(29)
    assert cached['AU5'].value == '1箱×4＋1箱×6'
    assert cached['P5'].value == 11 and cached['BC5'].value == 1.1
    with ZipFile(BytesIO(content)) as archive, ZipFile(TEMPLATE) as original:
        old_styles = ET.fromstring(original.read('xl/styles.xml'))
        new_styles = ET.fromstring(archive.read('xl/styles.xml'))
        for section in ('fonts', 'fills', 'borders', 'numFmts'):
            assert ET.tostring(old_styles.find('s:'+section, NS)) == ET.tostring(new_styles.find('s:'+section, NS))
        drawing = ET.fromstring(archive.read(DRAWING_PART))
        assert len(drawing.findall('.//xdr:pic', NS)) == 4
        assert len(drawing.findall('.//xdr:sp', NS)) == 0


def test_same_material_same_carton_multiple_sources_keep_one_loading_and_each_price():
    data = frozen({'a': source('A'), 'b': source('A', price='.3', item='ITEM-B', contract='CONTRACT-B', po='PO-B')},
                  [package('共箱', [('a', 3, 0), ('b', 2, 1)])])
    _, cached, formulas = output(data)
    assert cached['AZ4'].value == 'A' and cached['AZ5'].value is None
    assert cached['AU4'].value == '1箱×6' and cached['AT4'].value == 1
    assert cached['P4'].value == 7 and cached['AP4'].value == .9 and cached['AQ4'].value == .9
    assert cached['D4'].value == 'ITEM-A／ITEM-B' and cached['AN4'].value == 'PO-A／PO-B'
    assert 'CONTRACT-A／CONTRACT-B' in cached['BE4'].value
    assert formulas['AO4'].value == '=ROUND(AP4/K4,4)'


def test_same_material_multiple_suppliers_do_not_assign_row_total_to_each_supplier():
    data = frozen({'a': source('A'), 'b': source('A', price='.3', supplier='供应商乙', po='PO-B'), 'c': source('C')}, [
        package('共箱', [('a', 6, 0), ('b', 4, 0)]), package('C单装', [('c', 2, 0)])])
    content, cached, _ = output(data)
    assert cached['AL4'].value == '供应商甲／供应商乙'
    assert cached['AP4'].value == 1.8 and cached['AP5'].value == .2
    assert cached['AQ4'].value == '供应商甲: 0.80\n供应商乙: 1.20'
    assert cached['AQ4'].alignment.wrap_text
    assert cached['AQ5'].value is None
    workbook = detail_book(BytesIO(content), data_only=True)
    assert workbook[CALC_NAME].sheet_state == 'hidden'
    original = detail_book(TEMPLATE)[TEMPLATE_SHEET_NAME]
    assert copy(cached['AQ4'].font) == copy(original['AQ4'].font)
    assert copy(cached['AQ4'].border) == copy(original['AQ4'].border)
    assert cached['AQ4'].number_format == original['AQ4'].number_format


@pytest.mark.parametrize('item_count', [1, 3])
def test_clean_export_changes_only_detail_and_keeps_other_sheets(item_count):
    original_bytes = TEMPLATE.read_bytes()
    sources = {str(i): source(f'TEST-{i}') for i in range(item_count)}
    data = frozen(sources, [package(f'TEST箱{i}', [(key, 10, 1)]) for i, key in enumerate(sources)])
    content, cached, formulas = output(data)
    last_row = item_count + 3
    assert TEMPLATE.read_bytes() == original_bytes
    assert cached.title == SHEET_NAME
    assert cached.max_row == last_row and cached.max_column == 57
    assert str(cached.print_area) == f"'{SHEET_NAME}'!$A$3:$BE${last_row}"
    assert cached.auto_filter.ref == f'A3:BE{last_row}'
    original = detail_book(BytesIO(original_bytes))[TEMPLATE_SHEET_NAME]
    assert [cell.value for cell in cached[3]] == [cell.value for cell in original[3][:57]]
    for column in range(1, 58):
        actual, before = cached.cell(4, column), original.cell(4, column)
        assert copy(actual.font) == copy(before.font)
        assert copy(actual.fill) == copy(before.fill)
        assert copy(actual.border) == copy(before.border)
        assert actual.number_format == before.number_format
        assert actual.comment is None
    for key, dimension in cached.column_dimensions.items():
        assert dimension.width == original.column_dimensions[key].width
    assert formulas['P4'].value.startswith('=Q4+')
    assert cached['P4'].value == 12 and cached['AP4'].value == 1
    assert '备品 1 PCS，免费' in cached['BE4'].value
    with ZipFile(BytesIO(content)) as archive, ZipFile(BytesIO(original_bytes)) as source_zip:
        paths = archive.namelist()
        assert 'xl/comments2.xml' not in paths and 'xl/drawings/vmlDrawing2.vml' not in paths
        sheet = ET.fromstring(archive.read(SHEET_PART))
        assert [int(row.get('r')) for row in sheet.findall('s:sheetData/s:row', NS)] == list(range(3, last_row+1))
        assert sheet.find('s:dimension', NS).get('ref') == f'A3:BE{last_row}'
        drawing = ET.fromstring(archive.read(DRAWING_PART))
        assert len(drawing.findall('.//xdr:pic', NS)) == item_count
        assert not drawing.findall('.//xdr:sp', NS)
        photos = [p for p in paths if p.startswith('xl/media/indonesia-material-')]
        assert len(photos) == item_count
        expected_photo = base64.b64decode(sources['0']['material']['image_data'].split(',')[1])
        assert all(archive.read(p) == expected_photo for p in photos)
        old_book = ET.fromstring(source_zip.read('xl/workbook.xml'))
        new_book = ET.fromstring(archive.read('xl/workbook.xml'))
        old_tabs = old_book.find('s:sheets', NS)
        new_tabs = new_book.find('s:sheets', NS)
        assert [dict(tab.attrib) for tab in new_tabs[:-1]] == [dict(tab.attrib) for tab in old_tabs]
        assert new_tabs[-1].get('name') == CALC_NAME and new_tabs[-1].get('state') == 'hidden'
        detail_index = str(next(i for i, tab in enumerate(old_tabs) if tab.get('name') == SHEET_NAME))
        def other_names(book):
            return [(dict(n.attrib), n.text) for n in book.findall('s:definedNames/s:definedName', NS)
                    if n.get('localSheetId') != detail_index]
        assert other_names(new_book) == other_names(old_book)
        old_shared = ET.fromstring(source_zip.read('xl/sharedStrings.xml'))
        new_shared = ET.fromstring(archive.read('xl/sharedStrings.xml'))
        protected_strings = set()
        for path in source_zip.namelist():
            if path.startswith('xl/worksheets/') and path.endswith('.xml') and path != SHEET_PART:
                assert archive.read(path) == source_zip.read(path)
                root = ET.fromstring(source_zip.read(path))
                protected_strings.update(int(cell.find('s:v', NS).text) for cell in root.findall('.//s:c', NS) if cell.get('t') == 's')
        assert len(old_shared) == len(new_shared)
        assert all(ET.tostring(old_shared[i]) == ET.tostring(new_shared[i]) for i in protected_strings)
        # Preserve other sheets' images, notes, printer settings and WPS objects,
        # including dependencies outside the worksheets directory.
        detail_resources, pending = set(), [SHEET_PART]
        while pending:
            part = pending.pop()
            rel_path = relationship_path(part)
            if rel_path not in source_zip.namelist():
                continue
            detail_resources.add(rel_path)
            for rel in ET.fromstring(source_zip.read(rel_path)):
                if rel.get('TargetMode') == 'External':
                    continue
                target = target_path(part, rel.get('Target'))
                if target not in detail_resources:
                    detail_resources.add(target); pending.append(target)
        mutable = {SHEET_PART, 'xl/workbook.xml', 'xl/_rels/workbook.xml.rels',
                   'xl/styles.xml', 'xl/sharedStrings.xml', '[Content_Types].xml'} | detail_resources
        for path in source_zip.namelist():
            if path not in mutable and not path.endswith('/'):
                assert archive.read(path) == source_zip.read(path), path


def test_template_install_preserves_complete_workbook(tmp_path):
    installed = tmp_path/'complete-template.xlsx'
    prepare_template(TEMPLATE, installed)
    assert installed.read_bytes() == TEMPLATE.read_bytes()
    with ZipFile(installed) as archive:
        sheets = ET.fromstring(archive.read('xl/workbook.xml')).find('s:sheets', NS)
        assert [s.get('name') for s in sheets] == ['已到ARRIVED', '类别金额', '实业合同', '实业发票', '装箱单',
            SHEET_NAME, '商品汇总表', '发票', '销售合同', '装箱单 (2)', '草稿大单-1', '司机资料', '单位对照', 'WpsReserved_CellImgList']


def test_long_contract_and_purchase_lists_wrap_and_expand_without_losing_text():
    entries = {'a': source('A', contract='TEST-CONTRACT-0001', po='TEST-PURCHASE-0001'),
               'b': source('A', contract='TEST-CONTRACT-0002', po='TEST-PURCHASE-0002')}
    data = frozen(entries, [package('TEST混箱', [('a', 600, 0), ('b', 590, 10)], cartons=2)])
    data['remarks'] = 'TEST非真实出货；两个合同的物料资料均保留。'
    content, cached, _ = output(data)
    original = detail_book(TEMPLATE)[TEMPLATE_SHEET_NAME]
    assert cached['BE4'].alignment.wrap_text
    assert original.row_dimensions[4].height < cached.row_dimensions[4].height <= 409.5
    assert 'TEST-CONTRACT-0001／TEST-CONTRACT-0002' in cached['BE4'].value
    assert cached['AN4'].value == 'TEST-PURCHASE-0001／TEST-PURCHASE-0002'
    assert cached['L4'].value == 1190 and cached['P4'].value == 1202
    assert cached.column_dimensions['BE'].width == original.column_dimensions['BE'].width
    assert len(cached._images) == 1
    # Repeated rows reuse the same wrapping style instead of growing styles per row.
    multi = frozen({'a': source('A'), 'b': source('B')}, [package('混箱', [('a', 1, 0), ('b', 1, 0)])])
    _, sheet, _ = output(multi)
    assert sheet['BE4'].style_id == sheet['BE5'].style_id


def test_amount_does_not_depend_on_splitting_the_same_material_into_packaging_groups():
    sources = {'a': source('A', price='.005')}
    one = output(frozen(sources, [package('整箱', [('a', 2, 0)], cartons=2)]))[1]
    split = output(frozen(sources, [package('箱1', [('a', 1, 0)]), package('箱2', [('a', 1, 0)])]))[1]
    assert one['AP4'].value == split['AP4'].value == .01
    assert one['AQ4'].value == split['AQ4'].value == .01


def test_cbm_sums_original_rounding_of_each_physical_group():
    data = frozen({'a': source('A')}, [package('组1', [('a', 3, 0)], cartons=3, dims=('10','10','10')),
                                      package('组2', [('a', 4, 0)], cartons=4, dims=('10','10','10'))])
    _, cached, formulas = output(data)
    assert cached['AT4'].value == 7 and cached['S4'].value == 0
    assert formulas['S4'].value.startswith("=SUM('出货计算'!S")
    assert formulas['R4'].value == '=AW4*AX4*AY4/28316.75*0.0283'


def test_per_carton_formula_text_cache_uses_excel_half_up_rounding():
    data = frozen({'a': source('A', net='.0000005')}, [package('箱1', [('a', 1, 0)], weight='0'),
                                                    package('箱2', [('a', 3, 0)], weight='0')])
    _, cached, formulas = output(data)
    assert cached['BB4'].value == '1箱×0.000001 kg＋1箱×0.000002 kg'
    assert formulas['BB4'].value.startswith("='出货计算'!Y")


def test_valid_upper_bound_weights_prices_and_dimensions_do_not_overflow_decimal_context():
    from decimal import getcontext
    original_precision = getcontext().prec
    entry = source('A', price='1000000000000', net='1000000000000')
    entry['supply']['billing_factor'] = '.00000001'
    data = frozen({'a': entry}, [package('大数边界', [('a', '1000000000000', 0)],
                        weight='1000000000000', dims=('1000000000000',)*3)])
    _, cached, _ = output(data)
    assert cached['AP4'].value == pytest.approx(1e32)
    assert cached['Q4'].value == pytest.approx(1e24)
    assert cached['P4'].value == pytest.approx(1e24+1e12)
    assert getcontext().prec == original_precision


@pytest.mark.parametrize('conflict', ['dimension', 'weight', 'unit', 'supplier_currency', 'material_currency'])
def test_conflicting_single_row_inputs_are_rejected_without_guessing(conflict):
    sources = {'a': source('A'), 'b': source('A')}
    groups = [package('组1', [('a', 1, 0)]), package('组2', [('b', 1, 0)])]
    if conflict == 'dimension': groups[1]['length_cm'] = '28'
    if conflict == 'weight': sources['b']['material']['net_weight_kg'] = '2'
    if conflict == 'unit': sources['b']['material']['unit'] = 'SET'
    if conflict == 'supplier_currency': sources['b']['supply']['currency'] = 'CNY'
    if conflict == 'material_currency': sources['b']['supply'].update(supplier='供应商乙', currency='CNY')
    with pytest.raises(HTTPException) as error:
        export_shipment(frozen(sources, groups))
    assert error.value.status_code == 422
    assert '尚待' not in error.value.detail
