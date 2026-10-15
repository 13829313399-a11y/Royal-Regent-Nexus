"""All stock entry channels reconcile once, with distinct counterparties/units."""
import importlib
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
from test_material_handover import client, issuer, confirmed, posting, receipt, workspace, BASE, FACTORY
from test_warehouse_operations import payload, post
from test_fabric_procurement import book, row, stored
from test_fabric_procurement_tracking import save
from test_fabric_receiving import request, receive

MONTHLY = '/api/warehouse-operations/fabric/monthly'


def report(client, month='2026-07', **filters):
    result = client.get(MONTHLY, params=dict(factory_id=FACTORY, month=month, **filters))
    assert result.status_code == 200, result.text
    assert result.headers['cache-control'] == 'no-store'
    return result.json()


def test_all_channels_and_original_split_receipts_are_counted_once_and_exported(client):
    party, headers = issuer(client)
    manual = post(client, payload(client, 'RECEIPT', item_code='01-01-00001', quantity='20',
                                 counterparty=party['name'], source_no='=SUM(1,1)', lot='LOT-1'))
    doc = confirmed(client, headers)
    scan = client.post(BASE + '/documents/' + doc['id'] + '/post', json=posting(doc, 1, [receipt(amount='15')]))
    assert scan.status_code == 200, scan.text
    # Purchasing-source physical receipt has two rolls/locations, but one header quantity.
    save(client, book([row(供应商=party['name'], 物料编码='01-01-00001', 单价=0, 订单日期='2026/7/18')]))
    source = stored(client)['items'][0]
    result = receive(client, source, request(source, receipt_date='2026-07-18'))
    assert result.status_code == 200, result.text
    first = post(client, payload(client, 'ISSUE', batch_id=manual['id'], quantity='5', counterparty='车间一'))
    _, request_headers = issuer(client, '车间一', 'REQUESTOR')
    issue_doc = confirmed(client, request_headers, source='ISSUE-SCAN-1')
    issue_body = posting(issue_doc, workspace(client)['revision'], [dict(line_id='L1', quantity='3', batch_id=manual['id'])])
    result = client.post(BASE + '/documents/' + issue_doc['id'] + '/post', json=issue_body)
    assert result.status_code == 200, result.text
    post(client, payload(client, 'RETURN', original_id=first['id'], quantity='1'))
    post(client, payload(client, 'TRANSFER', batch_id=manual['id'], location='A02', quantity='2'))
    post(client, payload(client, 'REQUEST', quantity='100', counterparty='车间一'))
    post(client, payload(client, 'RECEIPT', item_code='01-01-00001', unit='Pcs', material_category='ACCESSORY', quantity='2', counterparty='临时供应商'))
    before = workspace(client)
    result = report(client)
    total = next(r for r in result['totals'] if r['unit'] == '码')
    assert Decimal(total['net_received']) == 85
    assert Decimal(total['manual_received']) == 70 and Decimal(total['scan_received']) == 15
    assert total['net_issued'] == '8' and total['returned'] == '1' and total['net_outbound'] == '7'
    assert total['manual_outbound'] == '4' and total['scan_outbound'] == '3'
    assert Decimal(total['stock_change']) == 78
    assert next(r for r in result['totals'] if r['unit'] == 'Pcs')['net_received'] == '2'
    suppliers = [r for r in result['summary'] if r['direction'] == 'IN' and r['unit'] == '码']
    assert len(suppliers) == 1 and suppliers[0]['party_key'] == 'supplier:' + party['supplier_id']
    assert suppliers[0]['document_count'] == 3  # manual + scan + original header
    assert len(result['details']) == 7 and result['cutting_interface'] == 'PENDING'
    procurement = next(r for r in result['details'] if r['origin'] == 'PROCUREMENT_RECEIVING')
    assert procurement['unit_price'] == '0' and procurement['price_unit'] == '' and procurement['currency'] == ''
    assert workspace(client) == before  # a report does not write or advance a revision
    incoming = report(client, direction='IN', party_key='supplier:' + party['supplier_id'])
    assert Decimal(incoming['totals'][0]['net_received']) == 85
    scan_only = report(client, direction='IN', channel='SCAN')
    assert scan_only['totals'][0]['net_received'] == '15'
    outgoing = report(client, direction='OUT')
    assert outgoing['summary'][0]['party_name'] == '车间一' and outgoing['totals'][0]['net_outbound'] == '7'
    assert not outgoing['closed'] and not outgoing['amounts_calculated']
    exported = client.get(MONTHLY + '/export', params=dict(factory_id=FACTORY, month='2026-07', version=result['version']))
    assert exported.status_code == 200, exported.text
    sheets = load_workbook(BytesIO(exported.content), data_only=False)
    source_cell = next(cell for cells in sheets['收发原单明细'] for cell in cells if cell.value == '=SUM(1,1)')
    assert source_cell.data_type == 's'
    assert any(cell.value == '01-01-00001' for cells in sheets['收发原单明细'] for cell in cells)
    assert len(sheets['收发原单明细']['A']) == 9
    # The internal name grouping never grants public ownership of manual vouchers.
    public = client.get(BASE + '/external/monthly', headers=headers, params={'month': '2026-07'}).json()
    assert public['totals'][0]['month_net_received'] == '15' and public['unlinked_receipts'] == []
    # Same-name master identities cannot steal old manual vouchers or silently merge.
    database = importlib.import_module('app.db')
    master_model = importlib.import_module('app.models.fabric_master').FabricMasterRecord
    with database.SessionLocal() as db:
        db.add(master_model(id='same-name-other', factory_id=FACTORY, kind='SUPPLIER', code='OTHER',
            name=party['name'], status='ACTIVE', revision=1, data_json='{}', updated_at='2026-07-18'))
        db.commit()
    ambiguous = report(client, direction='IN')
    assert len([r for r in ambiguous['summary'] if r['unit'] == '码']) == 2
    unassigned = next(r for r in ambiguous['summary'] if r['party_key'] == 'supplier-name:' + party['name'])
    assert Decimal(unassigned['net_received']) == 70
    assert Decimal(ambiguous['totals'][1]['net_received']) == 85


def test_cross_month_receipt_and_issue_reversals_preserve_prior_month_and_source_channel(client):
    original = post(client, payload(client, 'RECEIPT', quantity='20', counterparty='供方A'))
    issue = post(client, payload(client, 'ISSUE', batch_id=original['id'], quantity='5', counterparty='裁床'))
    spare = post(client, payload(client, 'RECEIPT', quantity='4', counterparty='供方B'))
    july = report(client)
    post(client, payload(client, 'CORRECTION', original_id=issue['id'], business_date='2026-08-01'))
    post(client, payload(client, 'CORRECTION', original_id=spare['id'], business_date='2026-08-02'))
    assert report(client) == july
    august = report(client, '2026-08')
    assert august['totals'][0]['net_received'] == '-4'
    assert august['totals'][0]['net_outbound'] == '-5'
    assert august['totals'][0]['stock_change'] == '1'
    assert august['totals'][0]['issue_reversed'] == '5' and august['totals'][0]['receipt_reversed'] == '4'
    assert all(r['correction'] and r['channel'] == 'MANUAL' for r in august['details'])
    assert report(client, '2026-08', party_key='recipient:裁床')['totals'][0]['net_outbound'] == '-5'
    assert report(client, direction='OUT', party_key='supplier-name:供方A')['details'] == []
    old_version = july['version']
    post(client, payload(client, 'RECEIPT', item_code='NEW', quantity='1', counterparty='供方C'))
    assert client.get(MONTHLY + '/export', params=dict(factory_id=FACTORY, month='2026-07', version=old_version)).status_code == 409


def test_internal_report_requires_scoped_read_and_never_accepts_supplier_token(client, monkeypatch):
    assert client.get(MONTHLY, params=dict(factory_id='huakang-d', month='2026-07')).status_code == 422
    assert client.get(MONTHLY, params=dict(factory_id=FACTORY, month='2026-13')).status_code == 422
    assert client.get(MONTHLY, params=dict(factory_id=FACTORY, month='2026-07', channel='CUTTING')).status_code == 422
    assert client.get(MONTHLY + '/export', params=dict(factory_id=FACTORY, month='2026-07', direction='OTHER')).status_code == 422
    api = importlib.import_module('app.api.warehouse_operations')
    with monkeypatch.context() as scoped_patch:
        scoped_patch.setattr(api, 'can', lambda *args, **kwargs: False)
        assert client.get(MONTHLY, params=dict(factory_id=FACTORY, month='2026-07')).status_code == 403
    _, headers = issuer(client)
    client.cookies.clear()
    assert client.get(MONTHLY, headers=headers, params=dict(factory_id=FACTORY, month='2026-07')).status_code == 401
    assert client.get(MONTHLY + '/export', headers=headers, params=dict(factory_id=FACTORY, month='2026-07')).status_code == 401


def test_search_unit_and_date_filters_aggregate_and_export_the_same_signed_evidence(client):
    original = post(client, payload(client, 'RECEIPT', quantity='0.100001', source_no='DN-0007', counterparty='供方甲'))
    post(client, payload(client, 'CORRECTION', original_id=original['id'], business_date='2026-07-20'))
    post(client, payload(client, 'RECEIPT', quantity='2', source_no='DN-0007', counterparty='供方甲',
                         item_code='PCS-1', unit='Pcs', material_category='ACCESSORY'))
    post(client, payload(client, 'RECEIPT', quantity='100', source_no='DN-OTHER', counterparty='供方乙'))
    before = workspace(client)
    filters = dict(search=' dn-0007 ', unit='码', date_order='DESC')
    result = report(client, **filters)
    assert len(result['summary']) == 1 and len(result['details']) == 2
    assert result['summary'][0]['document_count'] == 1
    assert Decimal(result['totals'][0]['net_received']) == 0
    assert Decimal(result['totals'][0]['received']) == Decimal('0.100001')
    assert result['details'][0]['correction'] and result['details'][0]['quantity'] == '-0.100001'
    assert result['units'] == ['Pcs', '码']
    assert len(result['parties']) == 2  # narrowing must not lose the option to switch party
    assert len(report(client, search=original['id'])['details']) == 2
    assert len(report(client, search='供方甲', unit='Pcs')['details']) == 1
    assert report(client, search='not-present')['totals'] == []
    assert report(client, search='dn-0007', unit='码', date_order='ASC')['details'][0]['id'] == original['id']
    exported = client.get(MONTHLY + '/export', params=dict(factory_id=FACTORY, month='2026-07', version=result['version'], **filters))
    assert exported.status_code == 200, exported.text
    sheets = load_workbook(BytesIO(exported.content), data_only=False)
    assert sheets['收发原单明细'].max_row == 4
    assert sheets['收发数量汇总'].max_row == 3
    assert workspace(client) == before
    mismatch = client.get(MONTHLY + '/export', params=dict(factory_id=FACTORY, month='2026-07', version=result['version'], search='DN-OTHER'))
    assert mismatch.status_code == 409
    for query in ({'date_order': 'INVALID'}, {'search': 'x' * 129}, {'unit': 'x' * 33}):
        assert client.get(MONTHLY, params=dict(factory_id=FACTORY, month='2026-07', **query)).status_code == 422
