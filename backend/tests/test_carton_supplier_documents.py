"""Supplier document visibility, batch acceptance and Excel export boundaries."""
from io import BytesIO
import json
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from pypdf import PdfWriter
from test_molding_sample_api import make_client, login_as
from test_carton_procurement_api import _create_order, _order_payload
from test_carton_supplier_portal import BASE, setup_portal
from test_carton_replenishment_api import setup as setup_replenishment, body as replenishment_body, replenish
from app.services.carton_supplier_portal import SUPPLIER_ORDER_IMPORT_HEADERS, _supplier_import_dimensions


def test_supplier_evidence_prevents_reassignment_after_attachment(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        order = _create_order(client, submit_supplier=False)
        assert order["status"] == "CONFIRMED"
        pdf = PdfWriter()
        pdf.add_blank_page(width=100, height=100)
        content = BytesIO()
        pdf.write(content)
        uploaded = client.post(BASE + f"/internal/orders/{order['id']}/attachments",
                               data={"factory_id": "huaxing"},
                               files={"file": ("supplier-a.pdf", content.getvalue())})
        assert uploaded.status_code == 201, uploaded.text
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonSupplier
        with SessionLocal() as db:
            db.add(CartonSupplier(id="new-supplier", factory_id="huaxing", supplier_code="OTHER",
                                  supplier_name="其他供应商", status="ACTIVE",
                                  created_at="2026-09-21", updated_at="2026-09-21"))
            db.commit()
        update = {**_order_payload(), "expected_revision": order["revision"],
                  "reason": "供应商改派安全回归", "supplier_id": "new-supplier"}
        response = client.patch(f"/api/carton-procurement/orders/{order['order_no']}", json=update)
        assert response.status_code == 409, response.text
        assert "不能改派供应商" in response.text


def test_batch_accept_is_atomic_per_factory_and_shipping_needs_acceptance(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client)
        order = client.get(BASE + "/workspace", params={"factory_id": "huaxing"}).json()["orders"][0]
        lines = [{
            "order_line_id": line["id"], "issue_id": order["issue_id"],
            "expected_revision": line["commitment_revision"], "promised_date": "2026-09-25",
        } for line in order["lines"]]
        assert len(lines) >= 2
        shipment = {
            "factory_id": "huaxing", "request_id": str(uuid4()),
            "delivery_note_no": "DN-BATCH-GATE", "delivery_date": "2026-09-25",
            "lines": [{"order_line_id": order["lines"][0]["id"], "issue_id": order["issue_id"], "quantity": 1}],
        }
        assert client.post(BASE + "/shipments", json=shipment).status_code == 409
        bad = [dict(line) for line in lines]
        bad[-1]["expected_revision"] += 1
        rejected = client.put(BASE + "/commitments/batch", json={"factory_id": "huaxing", "lines": bad})
        assert rejected.status_code == 409, rejected.text
        after_reject = client.get(BASE + "/workspace", params={"factory_id": "huaxing"}).json()["orders"][0]
        assert all(not line["accepted"] for line in after_reject["lines"])
        first = client.put(BASE + f"/papers/{lines[0]['order_line_id']}/commitment", json={
            "factory_id": "huaxing", "issue_id": lines[0]["issue_id"],
            "expected_revision": 0, "promised_date": "2026-09-25"})
        assert first.status_code == 200, first.text
        partly_shipped = client.post(BASE + "/shipments", json=shipment)
        assert partly_shipped.status_code == 409
        assert "整单接单" in partly_shipped.text
        accepted = client.put(BASE + "/commitments/batch", json={"factory_id": "huaxing", "lines": lines[1:]})
        assert accepted.status_code == 200, accepted.text
        assert accepted.json()["order_ids"] == [order["id"]]
        assert all(line["accepted"] for line in client.get(
            BASE + "/workspace", params={"factory_id": "huaxing"}).json()["orders"][0]["lines"])
        assert client.post(BASE + "/shipments", json=shipment).status_code == 201


def test_supplier_documents_export_and_audit_are_scoped(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client)
        docs = client.get(BASE + "/documents", params={"factory_id": "huaxing"})
        assert docs.status_code == 200, docs.text
        purchase = next(row for row in docs.json() if row["kind"] == "PURCHASE")
        assert purchase["export_count"] == 0
        # Counts must not include an event with the same document ID in another factory.
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonAuditEvent
        with SessionLocal() as db:
            db.add(CartonAuditEvent(
                id=f"CAE-{uuid4().hex}", factory_id="huakang-a",
                event_type="SUPPLIER_DOCUMENT_EXPORTED", entity_type="carton_purchase_order_issue",
                entity_id=purchase["id"], actor_user_id="test", actor_name="test",
                created_at="2026-09-24T00:00:00"))
            db.commit()
        purchase_after_other_factory_event = next(row for row in client.get(
            BASE + "/documents", params={"factory_id": "huaxing"}).json() if row["id"] == purchase["id"])
        assert purchase_after_other_factory_event["export_count"] == 0
        assert purchase["orders"][0]["contract_no"]
        assert purchase["lines"]
        assert "unit_price" not in docs.text and "supplier_id" not in docs.text
        assert client.get(BASE + "/documents", params={"factory_id": "huadeng"}).status_code in (403, 422)
        assert client.get(BASE + "/activity", params={"factory_id": "huadeng"}).status_code in (403, 422)

        order = client.get(BASE + "/workspace", params={"factory_id": "huaxing"}).json()["orders"][0]
        lines = [{
            "order_line_id": line["id"], "issue_id": order["issue_id"],
            "expected_revision": line["commitment_revision"], "promised_date": "2026-09-25",
        } for line in order["lines"]]
        assert client.put(BASE + "/commitments/batch", json={"factory_id": "huaxing", "lines": lines}).status_code == 200
        sent = client.post(BASE + "/shipments", json={
            "factory_id": "huaxing", "request_id": str(uuid4()),
            "delivery_note_no": "DN-DOCS", "delivery_date": "2026-09-25",
            "lines": [{"order_line_id": order["lines"][0]["id"], "issue_id": order["issue_id"], "quantity": 3}],
        })
        assert sent.status_code == 201, sent.text
        delivery = next(row for row in client.get(
            BASE + "/documents", params={"factory_id": "huaxing"}).json() if row["kind"] == "DELIVERY")
        assert delivery["export_count"] == 0
        assert delivery["document_no"] == "DN-DOCS"
        assert {row["order_no"] for row in delivery["orders"]} == {order["order_no"]}
        selected = [
            {"factory_id": "huaxing", "kind": "PURCHASE", "id": purchase["id"]},
            {"factory_id": "huaxing", "kind": "DELIVERY", "id": delivery["id"]},
        ]
        exported = client.post(BASE + "/documents/export.xlsx", json={"documents": selected})
        assert exported.status_code == 200, exported.text
        book = load_workbook(BytesIO(exported.content), read_only=True)
        assert book.sheetnames == ["采购单", "送货单"]
        purchase_rows = list(book["采购单"].values)
        delivery_rows = list(book["送货单"].values)
        assert purchase["document_no"] in [row[1] for row in purchase_rows[1:]]
        assert "DN-DOCS" in [row[1] for row in delivery_rows[1:]]
        assert isinstance(delivery_rows[1][13], (int, float))
        book.close()
        counted = {row["id"]: row["export_count"] for row in client.get(
            BASE + "/documents", params={"factory_id": "huaxing"}).json()}
        assert counted[purchase["id"]] == 1
        assert counted[delivery["id"]] == 1
        # Duplicate selections are rejected before generating a workbook or recording exports.
        assert client.post(BASE + "/documents/export.xlsx", json={"documents": [
            selected[1], selected[1],
        ]}).status_code == 422
        assert client.post(BASE + "/documents/export.xlsx", json={"documents": [
            selected[1],
        ]}).status_code == 200
        counted = {row["id"]: row["export_count"] for row in client.get(
            BASE + "/documents", params={"factory_id": "huaxing"}).json()}
        assert counted[purchase["id"]] == 1
        assert counted[delivery["id"]] == 2
        activity = client.get(BASE + "/activity", params={"factory_id": "huaxing"})
        assert activity.status_code == 200
        actions = {row["action"] for row in activity.json()}
        assert {"采购单已发行", "纸品已确认接单", "供应商已确认发货"}.issubset(actions)
        assert "detail_json" not in activity.text and "unit_price" not in activity.text
        assert client.post(BASE + "/documents/export.xlsx", json={"documents": [
            *selected, {"factory_id": "huadeng", "kind": "PURCHASE", "id": purchase["id"]},
        ]}).status_code in (403, 422)
        assert client.post(BASE + "/documents/export.xlsx", json={"documents": [
            {"factory_id": "huaxing", "kind": "DELIVERY", "id": "forged-id"},
        ]}).status_code == 404
        counted = {row["id"]: row["export_count"] for row in client.get(
            BASE + "/documents", params={"factory_id": "huaxing"}).json()}
        assert counted[purchase["id"]] == 1
        assert counted[delivery["id"]] == 2


def test_supplier_order_import_template_uses_issued_delta_and_business_keys(monkeypatch):
    assert _supplier_import_dimensions("18*12.5*17.25", "PO") == [18.0, 12.5, 17.25]
    with make_client(monkeypatch) as client:
        setup_portal(client)
        purchase = next(row for row in client.get(BASE + "/documents", params={"factory_id": "huaxing"}).json()
                        if row["kind"] == "PURCHASE")
        selection = {"factory_id": "huaxing", "kind": "PURCHASE", "id": purchase["id"]}
        exported = client.post(BASE + "/documents/order-import.xlsx", json={"documents": [selection]})
        assert exported.status_code == 200, exported.text
        book = load_workbook(BytesIO(exported.content), read_only=True, data_only=True)
        assert book.sheetnames == ["006 (2)"]
        rows = list(book.active.values)
        assert list(rows[0]) == SUPPLIER_ORDER_IMPORT_HEADERS
        assert len(rows) == 3
        assert all(row[0] == purchase["orders"][0]["contract_no"] for row in rows[1:])
        assert all(row[1] == purchase["orders"][0]["item_no"] for row in rows[1:])
        assert rows[1][3:6] == (31.5, 11.125, 11.25)
        assert rows[1][6] == float(purchase["lines"][0]["change_quantity"])
        assert rows[1][9] == "inch"
        assert rows[1][11].date().isoformat() == purchase["orders"][0]["planned_date"]
        assert rows[2][5] is None  # flat paper has only two dimensions
        book.close()
        counted = next(row for row in client.get(BASE + "/documents", params={"factory_id": "huaxing"}).json()
                       if row["id"] == purchase["id"])
        assert counted["export_count"] == 1
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonPurchaseOrderIssue
        with SessionLocal() as db:
            issue = db.get(CartonPurchaseOrderIssue, purchase["id"])
            snapshot = json.loads(issue.snapshot_json)
            snapshot["lines"][0]["specification"] = "18*12.5*17.25"
            snapshot["lines"][0]["dimension_unit"] = ""
            issue.snapshot_json = json.dumps(snapshot, ensure_ascii=False)
            db.commit()
        cm_export = client.post(BASE + "/documents/order-import.xlsx", json={"documents": [selection]})
        assert cm_export.status_code == 200, cm_export.text
        cm_book = load_workbook(BytesIO(cm_export.content), read_only=True, data_only=True)
        cm_row = list(cm_book.active.values)[1]
        assert cm_row[3:6] == (18, 12.5, 17.25)
        assert cm_row[9] == "cm"
        cm_book.close()
        rejected = client.post(BASE + "/documents/order-import.xlsx", json={"documents": [
            {**selection, "id": "forged-id"},
        ]})
        assert rejected.status_code == 404
        rejected_kind = client.post(BASE + "/documents/order-import.xlsx", json={"documents": [
            {**selection, "kind": "DELIVERY"},
        ]})
        assert rejected_kind.status_code == 422
        with SessionLocal() as db:
            issue = db.get(CartonPurchaseOrderIssue, purchase["id"])
            snapshot = json.loads(issue.snapshot_json)
            snapshot["lines"][0]["specification"] = "13213"
            issue.snapshot_json = json.dumps(snapshot, ensure_ascii=False)
            db.commit()
        malformed = client.post(BASE + "/documents/order-import.xlsx", json={"documents": [selection]})
        assert malformed.status_code == 422
        assert "无法拆成" in malformed.json()["detail"]
        counted = next(row for row in client.get(BASE + "/documents", params={"factory_id": "huaxing"}).json()
                       if row["id"] == purchase["id"])
        assert counted["export_count"] == 2


@pytest.mark.parametrize("responsibility", ["OWN", "SUPPLIER"])
def test_supplier_order_import_exports_replenishment_delta_with_settlement_note(monkeypatch, responsibility):
    with make_client(monkeypatch) as client:
        order, stocks = setup_replenishment(client, monkeypatch)
        created = replenish(client, order, replenishment_body(order, stocks, responsibility, qty=2))
        assert created.status_code == 201, created.text
        issue = created.json()["issue"]
        setup_portal(client)
        documents = client.get(BASE + "/documents", params={"factory_id": "huaxing"}).json()
        purchase = next(row for row in documents if row["id"] == issue["id"])
        assert purchase["replenishment"] is True
        assert purchase["export_count"] == 0
        selection = {"factory_id": "huaxing", "kind": "PURCHASE", "id": issue["id"]}
        response = client.post(BASE + "/documents/order-import.xlsx", json={"documents": [selection]})
        assert response.status_code == 200, response.text
        book = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        rows = list(book.active.values)
        book.close()
        assert len(rows) == 2  # The other original paper line has no replacement quantity.
        exported = rows[1]
        assert exported[0:2] == (purchase["orders"][0]["contract_no"], purchase["orders"][0]["item_no"])
        assert exported[6] == 2  # B01 delta, not the original cumulative demand.
        assert exported[18] == issue["document_no"]
        assert f"补单 {issue['document_no']}" in exported[13]
        if responsibility == "SUPPLIER":
            assert exported[10] == 0
            assert "供应商责任，免费换补" in exported[13]
        else:
            assert exported[10] > 0
            assert "我方责任，按实收结算" in exported[13]
        refreshed = client.get(BASE + "/documents", params={"factory_id": "huaxing"}).json()
        assert next(row for row in refreshed if row["id"] == issue["id"])["export_count"] == 1
