from io import BytesIO

from openpyxl import Workbook, load_workbook
import pytest

from test_internal_quote_api import make_client, login, logout
from test_internal_quote_alternatives import create, fill


def test_generated_formula_view_handles_references_lazy_branches_and_errors():
    from app.services.internal_quote_workbook_view import WorkbookValues, serialize_workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "报价明细"
    for address, value in {"A1": "ABS", "A2": "C-ABS", "B1": 1.005, "B2": 2,
                           "C1": '=SUMIF(A1:A2,"*ABS*",B1:B2)', "C2": '=IFERROR(1/0,3)',
                           "C3": '=ROUND(B1,2)', "C4": '=IF(B2>0,B2,1/0)',
                           "C5": "=MAX(1,ROUND((B1+B2)/2,0))", "C6": "=UNKNOWN(2)"}.items():
        sheet[address] = value
    other = book.create_sheet("隐藏计算")
    other.sheet_state = "veryHidden"
    other["A1"] = 4
    sheet["C7"] = "='隐藏计算'!$A$1+SUM(B1:B2)"
    calc = WorkbookValues(book)
    assert calc.cell(sheet.title, "C1") == pytest.approx(3.005)
    assert calc.cell(sheet.title, "C2") == 3
    assert calc.cell(sheet.title, "C3") == 1.01
    assert calc.cell(sheet.title, "C4") == 2
    assert calc.cell(sheet.title, "C5") == 2
    assert calc.cell(sheet.title, "C7") == pytest.approx(7.005)
    sheets, errors = serialize_workbook(book)
    assert len(sheets) == 1
    assert errors == ["报价明细!C6"]


def test_output_table_roundtrip_preview_export_and_stale_freeze(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "table-owner", "admin", "*", "*")
        quote = fill(client, create(client, "TABLE"))
        root = f"/api/internal-quotes/{quote['id']}"
        source_book = Workbook()
        source_book.active.title = "原始附件不附带"
        buffer = BytesIO()
        source_book.save(buffer)
        uploaded = client.post(root + "/attachments", data={"department": "engineering"},
                               files={"file": ("original.xlsx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert uploaded.status_code == 201, uploaded.text
        response = client.get(root + "/output-layout")
        assert response.status_code == 200, response.text
        layout = response.json()
        assert layout["formula_errors"] == []
        assert [s["name"] for s in layout["sheets"]] == ["报价明细"]
        group = next(g for g in layout["groups"] if g["label"] == "螺丝")
        assert group["editable"]
        cells = [c for row in layout["sheets"][0]["rows"] for c in row["cells"]]
        assert any(c.get("binding", {}).get("group") == group["id"] for c in cells if c.get("binding"))
        body = {"base_token": layout["base_token"], "edits": [
            {"field_id": group["fields"][0]["id"], "value": "螺丝改名"},
            {"field_id": group["fields"][1]["id"], "value": "0.8"},
        ]}
        preview = client.post(root + "/output-layout/preview", json=body)
        assert preview.status_code == 200, preview.text
        assert preview.json()["formula_errors"] == []
        unchanged = client.get(root).json()
        assert next(s for s in unchanged["sections"] if s["department"] == "engineering")["payload"]["materials"][0]["item"] == "螺丝"
        saved = client.put(root + "/output-layout", json=body)
        assert saved.status_code == 200, saved.text
        assert saved.json()["quote"]["id"] == quote["id"]
        assert saved.json()["quote"]["version_label"] == quote["version_label"]
        current = client.get(root).json()
        row = next(s for s in current["sections"] if s["department"] == "engineering")["payload"]["materials"][0]
        assert row["item"] == "螺丝改名" and float(row["unit_price_rmb"]) == .8
        reloaded = client.get(root + "/output-layout").json()
        assert reloaded["factory_price_hkd"] == preview.json()["factory_price_hkd"]
        assert client.put(root + "/output-layout", json=body).status_code == 409
        bad = {"base_token": reloaded["base_token"], "edits": [{"field_id": "engineering:materials/0/quantity", "value": 99}]}
        assert client.put(root + "/output-layout", json=bad).status_code == 400
        issued = client.post(root + "/direct-issue", json={"revision": current["header_revision"]})
        assert issued.status_code == 200, issued.text
        assert issued.json()["export_manifest"]["spreadsheet_attachments"] == []
        assert issued.json()["export_manifest"]["source_spreadsheet_attachments"][0]["file_name"] == "original.xlsx"
        exported = client.get(root + f"/exports/{issued.json()['id']}/download")
        book = load_workbook(BytesIO(exported.content))
        assert "原始附件不附带" not in book.sheetnames
        assert any("螺丝改名" in str(c.value) for row in book["报价明细"] for c in row)
        assert len(client.get(root + "/attachments").json()) == 1
        frozen = client.get(root + "/output-layout").json()
        assert frozen["read_only"] and all(not g["editable"] for g in frozen["groups"])
        body["base_token"] = frozen["base_token"]
        assert client.put(root + "/output-layout", json=body).status_code == 409


@pytest.mark.parametrize("mode,count", [("ordinary", 2), ("product", 3)])
def test_output_view_evaluates_generated_multi_carton_and_justplay_layouts(monkeypatch, mode, count):
    import test_internal_quote_excel_formula_links as existing
    from app.services.internal_quote_workbook_view import serialize_workbook
    books = []
    def book():
        result = Workbook()
        books.append(result)
        return result
    monkeypatch.setattr(existing, "Workbook", book)
    existing.test_carton_export_uses_each_packing_quantity_and_counts_flat_cards_once(mode, "cm", "cm", count)
    _, errors = serialize_workbook(books[0])
    assert errors == []


def test_multiple_groups_currencies_atomic_failure_and_molding_dependency(monkeypatch):
    from test_internal_quote_api import create_payload, ALL_SECTION_CODES
    from test_internal_quote_p2_api import ENGINEERING_PAYLOAD, MOLDING_PAYLOAD
    with make_client(monkeypatch) as client:
        login(client, "table-mixed", "admin", "*", "*")
        data = create_payload(suffix="TABLE-MIXED", participating_sections=ALL_SECTION_CODES)
        data["workflow_mode"] = "direct_output"
        created = client.post('/api/internal-quotes', json=data)
        assert created.status_code == 201, created.text
        root = f"/api/internal-quotes/{created.json()['id']}"
        engineering = {**ENGINEERING_PAYLOAD, "materials": [
            {"item": "同名物料", "category": "auxiliary", "quantity": 2, "unit_price_source_currency": "HKD", "unit_price_hkd": 3, "unit_price_rmb": 2.55},
            {"item": "同名物料", "category": "hardware", "quantity": 1, "unit_price_rmb": 5, "import_batch_id": "keep-lineage"},
        ]}
        original_groups = [{"id": f"board-{i}", "name": f"电子报价{i}", "components": [
            {"item": "同名IC", "quantity": 1, "unit_price_rmb": i, "unit_price_hkd": 999, "import_batch_id": f"source-{i}"},
        ]} for i in (1, 2)]
        painting = {"rows": [{"name": "喷油", "operations": {"clamp": {"quantity": 2, "unit_price_hkd": 1}},
                               "cost_allocation": "split", "paint_cost_hkd": .4, "labor_cost_hkd": 1.6}]}
        result = client.put(root + '/sections/save-all', json={"sections": [
            {"section_code": code, "revision": 1, "payload": payload} for code, payload in {
                "engineering": engineering, "molding": MOLDING_PAYLOAD, "electronic": {"quote_groups": original_groups}, "painting": painting,
            }.items()]})
        assert result.status_code == 200, result.text
        response = client.get(root + '/output-layout')
        assert response.status_code == 200, response.text
        layout = response.json()
        assert [s['name'] for s in layout['sheets']] == ['报价明细', '电子明细1', '电子明细2']
        second_sheet = layout['sheets'][2]
        second_price = next(c for row in second_sheet['rows'] for c in row['cells'] if c['address'] == 'E4')
        assert second_price['binding']['group'] == 'electronic:quote_groups/1/components/0'
        second_hkd = next(c for row in second_sheet['rows'] for c in row['cells'] if c['address'] == 'F4')
        assert float(second_hkd['text']) == pytest.approx(2 / .85, abs=.0001)
        ids = {f['id'] for g in layout['groups'] for f in g['fields']}
        assert 'molding:injection_lines/0/item' not in ids
        assert 'engineering:materials/0/unit_price_hkd' in ids
        assert 'engineering:materials/0/unit_price_rmb' not in ids
        body = {'base_token': layout['base_token'], 'edits': [
            {'field_id': 'engineering:materials/0/unit_price_hkd', 'value': 4},
            {'field_id': 'engineering:molds/0/item', 'value': '主模改名'},
            {'field_id': 'electronic:quote_groups/1/components/0/unit_price_rmb', 'value': 7},
            {'field_id': 'painting:rows/0/operations/clamp/unit_price_hkd', 'value': 2},
        ]}
        failed = client.put(root + '/output-layout', json=body)
        assert failed.status_code == 400, failed.text
        assert '合计' in failed.text
        assert client.get(root + '/output-layout').json()['base_token'] == layout['base_token']
        body['edits'].append({'field_id': 'painting:rows/0/labor_cost_hkd', 'value': 3.6})
        preview = client.post(root + '/output-layout/preview', json=body)
        assert preview.status_code == 200, preview.text
        result = client.put(root + '/output-layout', json=body)
        assert result.status_code == 200, result.text
        sections = {s['department']: s for s in result.json()['quote']['sections']}
        assert sections['engineering']['payload']['materials'][1] == engineering['materials'][1]
        assert sections['engineering']['payload']['materials'][0]['quantity'] == 2
        assert sections['electronic']['payload']['quote_groups'][0] == original_groups[0]
        assert sections['electronic']['payload']['quote_groups'][1]['components'][0]['import_batch_id'] == 'source-2'
        assert sections['molding']['calculation_status'] == 'valid'
        assert sections['molding']['dependency_status'] == 'current'
        assert client.get(root + '/output-layout').json()['factory_price_hkd'] == preview.json()['factory_price_hkd']
        assert sections['molding']['payload']['injection_lines'][0]['item'] == '主模改名'
        current = client.get(root + '/output-layout').json()
        changed_hkd = next(c for row in current['sheets'][2]['rows'] for c in row['cells'] if c['address'] == 'F4')
        assert float(changed_hkd['text']) == pytest.approx(7 / .85, abs=.0001)
        logout(client)
        login(client, 'table-department', 'molding_clerk', 'molding')
        forbidden = client.put(root + '/output-layout', json={'base_token': current['base_token'], 'edits': [
            {'field_id': 'engineering:materials/0/unit_price_hkd', 'value': 99}]})
        assert forbidden.status_code == 403, forbidden.text
        logout(client)
        login(client, 'table-other-factory', 'sales_customer_owner', 'sales-business', 'huakang-a')
        assert client.get(root + '/output-layout').status_code == 403


def test_visible_detail_subtotal_matches_server_with_repeated_parts_and_standalone_hair(monkeypatch):
    from test_internal_quote_p2_api import MOLDING_PAYLOAD
    with make_client(monkeypatch) as client:
        login(client, "table-total", "admin", "*", "*")
        from test_internal_quote_api import create_payload, ALL_SECTION_CODES
        payload = create_payload(suffix="TABLE-TOTAL", participating_sections=ALL_SECTION_CODES)
        payload["workflow_mode"] = "direct_output"
        response = client.post("/api/internal-quotes", json=payload)
        assert response.status_code == 201, response.text
        quote = fill(client, response.json())
        root = f"/api/internal-quotes/{quote['id']}"
        sections = {s["department"]: s for s in quote["sections"]}
        response = client.put(root + "/sections/save-all", json={"sections": [
            {"section_code": code, "revision": sections[code]["revision"], "payload": data}
            for code, data in {
                "molding": MOLDING_PAYLOAD,
                "sewing": {"quote_mode": "quick", "quick_quotes": [{"doll_name": "车衣", "unit_price_hkd": 2.8}]},
                "hair": {"lines": [{"item": "车发", "craft": "车发", "weight_g": 3, "unit": "PCS", "unit_price_hkd": 1.6}]},
            }.items()]})
        assert response.status_code == 200, response.text
        assert all(s["calculation_status"] == "valid" for s in response.json()["sections"])
        layout = client.get(root + "/output-layout").json()
        assert layout["formula_errors"] == []
        rows = layout["sheets"][0]["rows"]
        markup_index = next(index for index, row in enumerate(rows) if any(c["text"] == "×" for c in row["cells"]))
        subtotal = rows[markup_index - 1]
        amount = next(c for c in subtotal["cells"] if c["address"].startswith("D"))
        assert float(amount["text"].replace(",", "")) == pytest.approx(float(layout["factory_price_hkd"]), abs=.01)
        assert not any(c["text"] == "车缝未分类成本" for row in rows for c in row["cells"])


def test_output_layout_respects_legacy_document_archive(monkeypatch):
    import importlib
    from test_internal_quote_alternatives import fork
    with make_client(monkeypatch) as client:
        login(client, "table-archive", "admin", "*", "*")
        original = fill(client, create(client, "TABLE-ARCHIVE"))
        variant = fork(client, original)
        root = f"/api/internal-quotes/{variant['id']}"
        before = client.get(root + "/output-layout").json()
        db_module = importlib.import_module("app.db")
        models = importlib.import_module("app.models.internal_quote")
        with db_module.SessionLocal() as db:
            db.get(models.InternalQuote, original['id']).status = 'archived'
            db.commit()
        layout = client.get(root + "/output-layout").json()
        assert layout['read_only'] and all(not g['editable'] for g in layout['groups'])
        assert layout['base_token'] != before['base_token']
        body = {'base_token': layout['base_token'], 'edits': [
            {'field_id': 'engineering:materials/0/unit_price_rmb', 'value': 9}]}
        assert client.post(root + '/output-layout/preview', json=body).status_code == 409
        assert client.put(root + '/output-layout', json=body).status_code == 409


def test_output_layout_omits_unsynchronized_molding_cost_and_never_persists_on_read(monkeypatch):
    from test_internal_quote_projection_lifecycle import legacy_quote, stored_molding
    with make_client(monkeypatch) as client:
        login(client, "table-projection", "admin", "*", "*")
        quote_id, _ = legacy_quote(client)
        root = f"/api/internal-quotes/{quote_id}"
        before = stored_molding(quote_id)
        response = client.get(root + '/output-layout')
        assert response.status_code == 200, response.text
        layout = response.json()
        assert layout['read_only']
        assert any(w.get('code') == 'engineering_parts_unsaved' for w in layout['warnings'])
        summary = client.get(root + '/summary').json()
        assert float(layout['factory_price_hkd']) == float(summary['factory_price_hkd'])
        assert stored_molding(quote_id) == before
        body = {'base_token': layout['base_token'], 'edits': [
            {'field_id': 'engineering:materials/0/unit_price_rmb', 'value': 9}]}
        assert client.post(root + '/output-layout/preview', json=body).status_code == 409
        assert client.put(root + '/output-layout', json=body).status_code == 409
        assert stored_molding(quote_id) == before
