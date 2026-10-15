import copy
import importlib
from io import BytesIO
from uuid import uuid4
import pytest
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.exc import DatabaseError
from test_material_handover import client,issuer,confirmed,upload,posting,receipt,workspace,BASE,FACTORY


def test_monthly_uses_actual_dates_carryovers_prices_and_cross_month_corrections(client):
    person,headers=issuer(client)
    doc=upload(client,headers);data=copy.deepcopy(doc['data'])
    data.update(source_no='PRICE-001',document_date='2026-07-18')
    data['lines'][0].update(unit_price='0',price_unit='码',currency='CNY',price_note='原单免费补货，仅数量对账')
    result=client.post(BASE+'/external/documents/'+doc['id']+'/confirm',headers=headers,
        json=dict(request_id=str(uuid4()),expected_revision=doc['revision'],acknowledge_warnings=True,data=data))
    assert result.status_code==200,result.text
    doc=result.json()
    assert workspace(client)['stock']==[]
    result=client.post(BASE+'/documents/'+doc['id']+'/post',json=posting(doc,0,[receipt(amount='15')]))
    assert result.status_code==200,result.text
    doc=result.json();first_operation=doc['movements'][0]['operation_id']
    body=posting(doc,1,[receipt(amount='25')]);body['business_date']='2026-08-01'
    assert client.post(BASE+'/documents/'+doc['id']+'/post',json=body).status_code==200
    result=client.post('/api/warehouse-operations/fabric/documents',json=dict(factory_id=FACTORY,request_id=str(uuid4()),
        expected_revision=2,kind='CORRECTION',business_date='2026-08-02',original_id=first_operation,reason='七月登记数量复核后冲回'))
    assert result.status_code==200,result.text
    july=client.get(BASE+'/external/monthly',headers=headers,params={'month':'2026-07'}).json()
    august=client.get(BASE+'/external/monthly',headers=headers,params={'month':'2026-08'}).json()
    assert july['details'][0]['month_net_received']=='15' and july['details'][0]['outstanding']=='25'
    row=august['details'][0]
    assert row['prior_received']=='15' and row['month_received']=='25' and row['month_reversed']=='15'
    assert row['month_net_received']=='10' and row['outstanding']=='15' and row['unit_price']=='0'
    assert row['price_note'].startswith('原单免费') and not row['notes']
    assert not august['amounts_calculated'] and not august['closed']
    assert not august['unlinked_receipts']


def test_supplier_master_link_rotation_keeps_history_and_audit_without_storing_keys(client):
    person,headers=issuer(client);doc=confirmed(client,headers)
    body=dict(factory_id=FACTORY,request_id=str(uuid4()),expected_revision=person['revision'],valid_days=365)
    result=client.post(BASE+'/issuers/'+person['id']+'/renew',json=body)
    assert result.status_code==200,result.text
    new=result.json();new_headers={'X-Handover-Key':new['token']}
    assert new['id']==person['id'] and new['supplier_id']==person['supplier_id']
    assert client.get(BASE+'/external/context',headers=headers).status_code==401
    assert client.get(BASE+'/external/documents/'+doc['id'],headers=new_headers).status_code==200
    replay=client.post(BASE+'/issuers/'+person['id']+'/renew',json=body).json()
    assert replay['replayed'] and 'token' not in replay
    records=client.get('/api/fabric-warehouse/master',params=dict(factory_id=FACTORY,kind='SUPPLIER')).json()['items']
    master=next(row for row in records if row['id']==person['supplier_id'])
    updated=client.post('/api/fabric-warehouse/master',json=dict(factory_id=FACTORY,request_id=str(uuid4()),id=master['id'],
        expected_revision=master['revision'],kind='SUPPLIER',code=master['code'],name='更名后的供应商',status='ACTIVE',data=master['data']))
    assert updated.status_code==200,updated.text
    context=client.get(BASE+'/external/context',headers=new_headers).json()
    assert context['name']=='更名后的供应商'
    old_doc=client.get(BASE+'/external/documents/'+doc['id'],headers=new_headers).json()
    assert old_doc['data']['counterparty']==person['name']
    database=importlib.import_module('app.db')
    with database.SessionLocal() as db:
        saved=db.execute(text('SELECT detail_json,result_json FROM material_handover_link_commands')).all()
        assert all(person['token'] not in str(row) and new['token'] not in str(row) for row in saved)
        for table in ('material_handover_supplier_bindings','material_handover_link_commands'):
            with pytest.raises(DatabaseError):db.execute(text(f'DELETE FROM {table}'))
            db.rollback()


def test_monthly_isolated_from_other_supplier_and_flags_non_linked_receipts(client):
    party,headers=issuer(client);_,other=issuer(client,'其他供方')
    doc=confirmed(client,headers)
    assert client.post(BASE+'/documents/'+doc['id']+'/post',json=posting(doc,0,[receipt(amount='5')])).status_code==200
    assert client.get(BASE+'/external/monthly',headers=other,params={'month':'2026-07'}).json()['details']==[]
    assert client.get(BASE+'/external/monthly',headers=other,params={'month':'2026-07','supplier_id':party['supplier_id']}).status_code==403
    body=dict(factory_id=FACTORY,request_id=str(uuid4()),expected_revision=1,kind='RECEIPT',business_date='2026-07-18',
        source_no='LEGACY-VOUCHER',source_line='1',item_code='OLD',item_name='旧入口布料',unit='码',lot='old',location_id='fabric-A01',quantity='3',counterparty='未关联供方')
    assert client.post('/api/warehouse-operations/fabric/documents',json=body).status_code==200
    internal=client.get(BASE+'/monthly',params={'factory_id':FACTORY,'month':'2026-07'}).json()
    assert len(internal['unlinked_receipts'])==1 and internal['totals'][0]['month_net_received']=='5'
    public=client.get(BASE+'/external/monthly',headers=headers,params={'month':'2026-07'}).json()
    assert public['unlinked_receipts']==[]
    assert client.get(BASE+'/monthly',params={'factory_id':FACTORY,'month':'2026-13'}).status_code==422
    exported=client.get(BASE+'/external/monthly/export',headers=headers,params={'month':'2026-07'})
    assert exported.status_code==200
    book=load_workbook(BytesIO(exported.content),data_only=False)
    assert book['单据数量对账']['D3'].value=='01-01-00001'
    assert book['未关联入口的库存凭证'].max_row==1


def test_confirmation_prices_are_optional_but_identity_and_values_are_checked(client):
    party,headers=issuer(client);doc=upload(client,headers);data=copy.deepcopy(doc['data'])
    data.update(source_no='DN-PRINT',document_date='2026-07-18');data['lines'][0]['quantity']='13.5'
    body=dict(request_id=str(uuid4()),expected_revision=doc['revision'],acknowledge_warnings=True,data=data)
    invalid=copy.deepcopy(body);invalid['data']['supplier_id']='foreign'
    assert client.post(BASE+'/external/documents/'+doc['id']+'/confirm',headers=headers,json=invalid).status_code==403
    invalid=copy.deepcopy(body);invalid['data']['lines'][0]['unit_price']='-1'
    assert client.post(BASE+'/external/documents/'+doc['id']+'/confirm',headers=headers,json=invalid).status_code==422
    result=client.post(BASE+'/external/documents/'+doc['id']+'/confirm',headers=headers,json=body)
    assert result.status_code==200 and result.json()['data']['lines'][0]['quantity']=='13.5'
    assert client.post(BASE+'/external/documents/'+doc['id']+'/confirm',headers=headers,json=body).json()==result.json()
    assert workspace(client)['stock']==[]


def test_legacy_link_requires_explicit_master_binding_and_keeps_original_documents(client):
    party,headers=issuer(client)
    database=importlib.import_module('app.db');model=importlib.import_module('app.models.material_handover').HandoverIssuer
    import hashlib
    token='x'*43
    with database.SessionLocal() as db:
        db.add(model(id='legacy',factory_id=FACTORY,name=party['name'],kind='SUPPLIER',token_hash=hashlib.sha256(token.encode()).hexdigest(),
            expires_at='2030-01-01',active=1,revision=1,created_by='fixture',created_at='2026-07-01'));db.commit()
    legacy_headers={'X-Handover-Key':token};doc=upload(client,legacy_headers)
    body=dict(factory_id=FACTORY,request_id=str(uuid4()),expected_revision=1,supplier_id=party['supplier_id'])
    result=client.post(BASE+'/issuers/legacy/bind',json=body)
    assert result.status_code==200,result.text
    assert client.get(BASE+'/external/documents/'+doc['id'],headers=headers).status_code==200
    _,other=issuer(client,'另一个供应商')
    assert client.get(BASE+'/external/documents/'+doc['id'],headers=other).status_code==404
