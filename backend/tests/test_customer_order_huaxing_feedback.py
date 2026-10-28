"""Huaxing feedback acceptance: optional fields, hard conflicts and new-row color."""
from copy import copy
from io import BytesIO

import openpyxl
import msoffcrypto
import pytest
from fastapi import HTTPException
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font

from app.api import customer_order as api
from app.services import customer_order_huaxing_unified as mapping
from app.services import customer_order_unified as service
from app.services.customer_order_buzzbee import CustomerOrderWorkbookError, _decrypt_schedule
from app.services.customer_order_manual import decorate_manual_resolution_policy
from test_customer_order_huaxing_unified import template, data
from test_customer_order_seasons_factory import arguments


def test_wrong_schedule_password_has_customer_neutral_error():
    """Exercise an actual encrypted workbook through shared and unified loaders."""
    encrypted = BytesIO()
    msoffcrypto.OfficeFile(BytesIO(arguments('huaxing')['schedule_content'])).encrypt(
        'not-the-configured-password', encrypted)
    for loader, error_type in ((_decrypt_schedule, CustomerOrderWorkbookError),
                               (service._load_workbook, service.CustomerOrderUnifiedError)):
        with pytest.raises(error_type, match='排期工作簿密码') as caught:
            loader(encrypted.getvalue())
        assert 'BuzzBee' not in str(caught.value)


def conflict_arguments():
    args = arguments('huaxing')
    book = openpyxl.load_workbook(BytesIO(args['schedule_content']))
    book['ITEM表']['D4'] = book['ITEM表']['E4'] = 'ZE618957'
    book['ITEM表']['F4'] = 'MPO1609793'
    args['schedule_content'] = data(book)  # History has 24; incoming PO has 120.
    return args


def test_quantity_conflict_cannot_be_exported_with_forged_confirmation_keys():
    args = conflict_arguments()
    preview = decorate_manual_resolution_policy(service.create_unified_customer_preview(**args))
    row, = preview['rows']
    conflict = next(i for i in row['issues'] if i['code'] == 'existing_quantity_conflict')
    assert conflict['severity'] == 'blocked'
    assert not conflict['can_skip'] and not conflict['can_edit']
    with pytest.raises(HTTPException):
        api._actual_confirmed_issue_keys(preview, {conflict['skip_key']})
    with pytest.raises(service.CustomerOrderUnifiedError):
        service.export_unified_customer_schedule(**args, skipped_issue_keys={conflict['skip_key']})


def test_batch_quantity_conflict_remains_hard_after_advisory_policy():
    rows = [dict(id=str(i), product_no='W90001', reference_no='QF90000001',
                 quantity=quantity, issues=[], _huaxing_review_customer='seasons')
            for i, quantity in enumerate(('864', '288'))]
    service._mark_batch_duplicates(rows)
    mapping.apply_advisory_policy(rows)
    decorate_manual_resolution_policy({'rows': rows})
    issue = next(i for i in rows[1]['issues'] if i['code'] == 'existing_quantity_conflict')
    assert issue['severity'] == 'blocked' and issue['requires_source_correction']
    assert not issue['can_skip'] and not issue['can_edit']


def test_optional_date_code_can_stay_blank_or_be_exported_as_text():
    args = arguments('huaxing')
    preview = decorate_manual_resolution_policy(service.create_unified_customer_preview(**args))
    row, = preview['rows']
    issue = next(i for i in row['issues'] if i['field'] == 'date_code')
    assert issue['severity'] == 'warning' and issue['can_edit']
    empty, _, _ = service.export_unified_customer_schedule(**args)
    assert openpyxl.load_workbook(BytesIO(empty))['ITEM表']['U5'].value is None
    override = dict(row_id=row['id'], issue_key=issue['skip_key'], field='date_code', value='00123APR27')
    assert api._actual_manual_overrides(preview, [override]) == [override]
    output, _, exported = service.export_unified_customer_schedule(
        **args, skipped_issue_keys={issue['skip_key']}, manual_overrides=[override])
    cell = openpyxl.load_workbook(BytesIO(output))['ITEM表']['U5']
    assert cell.value == '00123APR27' and cell.data_type == 's'
    assert exported['rows'][0]['date_code'] == '00123APR27'


def test_optional_date_code_override_passes_api_export_control(monkeypatch):
    args = arguments('huaxing')
    preview = decorate_manual_resolution_policy(service.create_unified_customer_preview(**args))
    row, = preview['rows']
    issue = next(i for i in row['issues'] if i['field'] == 'date_code')
    override = dict(row_id=row['id'], issue_key=issue['skip_key'], field='date_code', value='00123APR27')
    monkeypatch.setattr(api, '_ensure_preview_fingerprint', lambda *args: 'verified')
    monkeypatch.setattr(api, '_record_export_audit', lambda *args, **kwargs: 'audit')
    _, audit = api._complete_export_control(
        None, None, preview=preview, received_date=args['received_date'],
        submitted_fingerprint='verified', requested_issue_keys={issue['skip_key']},
        requested_manual_overrides=[override], confirmation_reason='已核对原单日期码',
        output=b'test', file_name='test.xlsx')
    assert audit == 'audit'


def test_dickie_colors_apply_only_to_new_rows_and_duplicate_red_takes_priority():
    book = template('dickie')
    for name in ('ITEM表', '接单表', '正单评审表'):
        sheet = book[name]
        sheet['H4' if name == 'ITEM表' else 'G4'] = 'OLD-ITEM'
        sheet['E4' if name == 'ITEM表' else 'D4'] = 'OLD-ORDER'
        sheet['A4'].font = Font(name='Arial', color='FF008000', bold=True)
        sheet.conditional_formatting.add('A4:S10', FormulaRule(formula=['FALSE'], font=Font(color='FF999999')))
    before = {name: [(cell.coordinate, cell.value, copy(cell.font)) for cell in book[name][4]]
              for name in ('ITEM表', '接单表', '正单评审表')}
    rows = [dict(id=str(index), _huaxing_customer='dickie', item_sheet_name='ITEM表',
                 reference_no=f'NEW-{index}', product_no=f'ITEM-{index}', quantity='120',
                 units_per_carton='12', issues=issues)
            for index, issues in enumerate(([], [dict(code='duplicate_order_warning')]))]
    mapping.export_rows(book, rows)
    book = openpyxl.load_workbook(BytesIO(data(book)))
    for name, old_cells in before.items():
        sheet = book[name]
        for coordinate, value, font in old_cells:
            assert sheet[coordinate].value == value and sheet[coordinate].font == font
        assert sheet['A5'].font.color.rgb == 'FF0000FF'
        assert sheet['A6'].font.color.rgb == 'FFFF0000'
        assert any(rule.formula == ['FALSE'] for rules in sheet.conditional_formatting._cf_rules.values() for rule in rules)
        generated_ranges = [str(key.sqref) for key, rules in sheet.conditional_formatting._cf_rules.items()
                            if any(rule.formula == ['TRUE'] for rule in rules)]
        assert len(generated_ranges) == 2
        assert all('4:' not in value and '10' not in value for value in generated_ranges)


def test_dickie_repairs_wrong_summary_duplicate_columns_and_broken_row_references_only():
    book = template('dickie')
    summary = book['正单评审表']
    wrong = 'AND($E9<>"",$H9<>"",COUNTIFS($E$4:$E$1995,$E9,$H$4:$H$1995,$H9)>1)'
    broken = 'AND(#REF!<>"",#REF!<>"",COUNTIFS($E$4:$E$1995,#REF!,$H$4:$H$1995,#REF!)>1)'
    summary.conditional_formatting.add('D9:G10', FormulaRule(formula=[wrong]))
    summary.conditional_formatting.add('D11:G11', FormulaRule(formula=[broken]))
    summary.conditional_formatting.add('N9:N10', FormulaRule(formula=['N9>100']))
    mapping._repair_dickie_duplicate_rules(book)
    formulas = {str(key.sqref): rules[0].formula[0] for key, rules in summary.conditional_formatting._cf_rules.items()}
    assert 'LEFT(UPPER(TRIM($D9)),2)="SC"' in formulas['D9:G10']
    assert 'COUNTIFS($D$4:' in formulas['D9:G10'] and '$G$4:' in formulas['D9:G10']
    assert 'SEARCH("MO",$B9)' in formulas['D9:G10']
    assert 'SEARCH("备料",$B9)' in formulas['D9:G10']
    assert '"<>*MO*"' in formulas['D9:G10'] and '"<>*备料*"' in formulas['D9:G10']
    assert '$E' not in formulas['D9:G10'] and '$H' not in formulas['D9:G10']
    assert '$D11' in formulas['D11:G11'] and '#REF!' not in formulas['D11:G11']
    assert formulas['N9:N10'] == 'N9>100'
