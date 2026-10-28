from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile
from urllib.parse import parse_qs, urlsplit

from lxml import etree as ET
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, text
from sqlalchemy.orm import Session
from sqlalchemy.exc import DatabaseError
from fastapi import HTTPException
import pytest

from test_online_documents import api, new, callback, open_editor, put_grants, BASE, SCOPE, context
from app.models.online_documents import OnlineDocument as Document, OnlineDocumentRevision as Revision, OnlineDocumentEvent as Event
from app.services import online_documents as service, online_document_management as management, online_document_protection as protection, online_document_protocol as protocol
from app.services.document_tools import storage


def payload(client, doc, **extra):
    doc_info = client.get(f"{BASE}/{doc}{SCOPE}").json()
    return {"request_id": uuid4().hex, "expected_version": doc_info["state_version"], **extra}


def upload(client):
    book = Workbook(); sheet = book.active
    sheet.title = "进出货"; sheet['A1'] = '产品'; sheet['B1'] = 7; sheet['B2'] = 8; sheet['C1'] = '=SUM(B1:B2)'
    stream = BytesIO(); book.save(stream)
    response = client.post(BASE + '/upload', data={'factory_id': 'huaxing'}, files={'file': ('表格.xlsx', stream.getvalue())})
    assert response.status_code == 201, response.text
    return response.json()['id']


def enable(client, doc):
    response = client.put(f"{BASE}/{doc}/protection{SCOPE}", json=payload(client, doc, enabled=True, rules=[{'sheet': '进出货', 'ranges': ['B1:C2']}]))
    assert response.status_code == 200, response.text
    return client.get(f"{BASE}/{doc}/download{SCOPE}").content


def mutate(data, part, edit):
    stream = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(stream, 'w') as output:
        for item in source.infolist():
            value = source.read(item.filename)
            if item.filename == part:
                root = ET.fromstring(value); edit(root); value = ET.tostring(root)
            output.writestr(item, value)
    return stream.getvalue()


def test_owner_history_restore_preserves_bytes_and_idempotency(api):
    client, engine, users, current, monkey = api
    doc = upload(client)
    original = client.get(f"{BASE}/{doc}/download{SCOPE}").content
    enabled = enable(client, doc)
    disabled = client.put(f"{BASE}/{doc}/protection{SCOPE}", json=payload(client, doc, enabled=False)).json()
    request = payload(client, doc, revision=1)
    result = client.post(f"{BASE}/{doc}/restore{SCOPE}", json=request)
    assert result.status_code == 200, result.text
    assert result.json()['revision'] == 4
    assert client.get(f"{BASE}/{doc}/download{SCOPE}").content == original
    assert client.post(f"{BASE}/{doc}/restore{SCOPE}", json=request).json()['revision'] == 4
    assert client.post(f"{BASE}/{doc}/restore{SCOPE}", json={**request, 'revision': 2}).status_code == 409
    assert client.post(f"{BASE}/{doc}/restore{SCOPE}", json={**request, 'request_id': uuid4().hex}).status_code == 409
    history = client.get(f"{BASE}/{doc}/history{SCOPE}").json()['items']
    assert [v['revision'] for v in history] == [4, 3, 2, 1]
    assert client.get(f"{BASE}/{doc}/history/2/download{SCOPE}").content == enabled
    with Session(engine) as db:
        assert db.scalar(select(Event).where(Event.command_id == request['request_id'])).details['restored_revision'] == 1
        assert db.query(Revision).count() == 4
        assert storage.resolve(db.scalars(select(Revision).order_by(Revision.revision)).first().storage_key).read_bytes() == original


@pytest.mark.parametrize('role', ['viewer', 'editor', 'stranger'])
def test_management_and_pregrant_history_owner_only(api, role):
    client, engine, users, current, monkey = api
    doc = upload(client)
    command = payload(client, doc, revision=1)
    put_grants(client, f"{BASE}/{doc}/grants{SCOPE}", json={'grants': [{'user_id': 'viewer', 'access': 'view'}, {'user_id': 'editor', 'access': 'edit'}]})
    current['user'] = role
    status = 404 if role == 'stranger' else 403
    for path in ('history', 'history/1/download', 'events', 'protection'):
        assert client.get(f"{BASE}/{doc}/{path}{SCOPE}").status_code == status
    for path in ('restore', 'trash', 'recover'):
        assert client.post(f"{BASE}/{doc}/{path}{SCOPE}", json=command).status_code == status
    assert client.put(f"{BASE}/{doc}/protection{SCOPE}", json={**command, 'enabled': True, 'rules': []}).status_code == status


def test_recycle_preserves_versions_revokes_links_and_restores_private(api):
    client, engine, users, current, monkey = api
    doc = upload(client)
    original = client.get(f"{BASE}/{doc}/download{SCOPE}").content
    put_grants(client, f"{BASE}/{doc}/grants{SCOPE}", json={'grants': [{'user_id': 'viewer', 'access': 'view'}]})
    current['user'] = 'viewer'; client.cookies.set('rr_session', 'viewer')
    config = open_editor(client, doc)
    token = parse_qs(urlsplit(config['document']['url']).query)['token'][0]
    current['user'] = 'owner'; client.cookies.set('rr_session', 'owner')
    request = payload(client, doc)
    deleted = client.post(f"{BASE}/{doc}/trash{SCOPE}", json=request).json()
    assert deleted['deleted_at'] and not deleted['grants']
    assert client.post(f"{BASE}/{doc}/trash{SCOPE}", json=request).status_code == 200
    assert client.get(BASE + SCOPE).json()['items'] == []
    assert len(client.get(BASE + SCOPE + '&deleted=true').json()['items']) == 1
    for path in ('', '/download', '/history', '/history/1/download'):
        assert client.get(f"{BASE}/{doc}{path}{SCOPE}").status_code == 404
    assert client.get(f"{BASE}/{doc}/source", params={'token': token}).status_code in {403, 404}
    assert callback(client, doc, {'status': 4, 'key': 'old'}).status_code == 403
    current['user'] = 'viewer'
    assert client.get(BASE + SCOPE + '&deleted=true').json()['items'] == []
    current['user'] = 'owner'
    request = {'request_id': uuid4().hex, 'expected_version': deleted['state_version']}
    restored = client.post(f"{BASE}/{doc}/recover{SCOPE}", json=request).json()
    assert not restored['deleted_at'] and not restored['grants'] and restored['revision'] == 1
    assert client.post(f"{BASE}/{doc}/recover{SCOPE}", json=request).status_code == 200
    assert client.get(f"{BASE}/{doc}/download{SCOPE}").content == original
    put_grants(client, f"{BASE}/{doc}/grants{SCOPE}", json={'grants': [{'user_id': 'viewer', 'access': 'view'}]})
    assert client.get(f"{BASE}/{doc}/source", params={'token': token}).status_code == 403


def test_active_sessions_block_all_management_and_epoch_fence(api):
    client, engine, users, current, monkey = api
    doc = upload(client); request = payload(client, doc, revision=1)
    open_editor(client, doc)
    for path in ('restore', 'trash'):
        assert client.post(f"{BASE}/{doc}/{path}{SCOPE}", json=request).status_code == 409
    assert client.put(f"{BASE}/{doc}/protection{SCOPE}", json={**request, 'enabled': True, 'rules': []}).status_code == 409
    users['owner'] = context('owner', epoch=2)
    assert client.get(f"{BASE}/{doc}/history{SCOPE}").status_code == 404


def test_protection_keeps_original_parts_and_valid_input_save(api):
    client, engine, users, current, monkey = api
    doc = upload(client)
    original = client.get(f"{BASE}/{doc}/download{SCOPE}").content
    protected = enable(client, doc)
    book = load_workbook(BytesIO(protected))
    assert book.active['C1'].protection.locked
    assert not book.active['B1'].protection.locked
    assert book.active['A1'].protection.locked
    assert book.active.protection.sheet and book.security.lockStructure
    with ZipFile(BytesIO(original)) as a, ZipFile(BytesIO(protected)) as b:
        for name in a.namelist():
            if name not in {'xl/workbook.xml', 'xl/styles.xml', 'xl/worksheets/sheet1.xml'}:
                assert a.read(name) == b.read(name)
    candidate = mutate(protected, 'xl/worksheets/sheet1.xml', lambda root: setattr(root.find('.//' + protection.Q('c') + '[@r="B1"]/' + protection.Q('v')), 'text', '15'))
    # The real editor omits openpyxl's empty definedNames container when saving.
    candidate = mutate(candidate, 'xl/workbook.xml', lambda root: root.remove(root.find(protection.Q('definedNames'))))
    protection.validate_saved(protected, candidate, [{'sheet': '进出货', 'ranges': ['B1:C2']}])
    key = open_editor(client, doc)['document']['key']
    monkey.setattr(protocol, 'fetch_saved', lambda url: candidate)
    assert callback(client, doc, {'status': 2, 'key': key, 'users': ['owner_1'], 'url': 'http://onlyoffice/cache/files/test'}).json() == {'error': 0}
    assert client.get(f"{BASE}/{doc}{SCOPE}").json()['revision'] == 3


@pytest.mark.parametrize('tamper', ['formula', 'locked', 'delete_formula', 'new_formula', 'sheet_name', 'unprotect', 'defined_name', 'date1904', 'merge'])
def test_protection_rejects_tampering_even_when_owner_is_last_editor(api, tamper):
    client, engine, users, current, monkey = api
    doc = upload(client); original = enable(client, doc)
    q = protection.Q
    def sheet_edit(root):
        if tamper == 'formula': root.find(f'.//{q("c")}[@r="C1"]/{q("f")}').text = '0'
        if tamper == 'locked': root.find(f'.//{q("c")}[@r="A1"]/{q("is")}/{q("t")}').text = 'tamper'
        if tamper == 'delete_formula':
            cell = root.find(f'.//{q("c")}[@r="C1"]'); cell.remove(cell.find(q('f')))
        if tamper == 'new_formula': ET.SubElement(root.find(f'.//{q("c")}[@r="B1"]'), q('f')).text = '1+1'
        if tamper == 'unprotect': root.remove(root.find(q('sheetProtection')))
        if tamper == 'merge': ET.SubElement(ET.SubElement(root, q('mergeCells'), count='1'), q('mergeCell'), ref='A1:B1')
    def book_edit(root):
        if tamper == 'sheet_name': root.find(q('sheets'))[0].set('name', '另一个表')
        if tamper == 'defined_name':
            names = root.find(q('definedNames'))
            if names is None: names = ET.SubElement(root, q('definedNames'))
            ET.SubElement(names, q('definedName'), name='Injected').text = '1'
        if tamper == 'date1904': root.find(q('workbookPr')).set('date1904', '1')
    candidate = mutate(original, 'xl/workbook.xml' if tamper in {'sheet_name', 'defined_name', 'date1904'} else 'xl/worksheets/sheet1.xml', book_edit if tamper in {'sheet_name', 'defined_name', 'date1904'} else sheet_edit)
    put_grants(client, f"{BASE}/{doc}/grants{SCOPE}", json={'grants': [{'user_id': 'editor', 'access': 'edit'}]})
    current['user'] = 'editor'; client.cookies.set('rr_session', 'editor'); key = open_editor(client, doc)['document']['key']
    current['user'] = 'owner'; client.cookies.set('rr_session', 'owner'); assert open_editor(client, doc)['document']['key'] == key
    monkey.setattr(protocol, 'command', lambda p: {'error': 0})
    request = client.post(f"{BASE}/{doc}/save{SCOPE}").json()['request_id']
    monkey.setattr(protocol, 'fetch_saved', lambda url: candidate)
    result = callback(client, doc, {'status': 6, 'forcesavetype': 0, 'userdata': request, 'key': key, 'users': ['owner_1'], 'url': 'http://onlyoffice/cache/files/test'})
    assert result.json() == {'error': 1}
    info = client.get(f"{BASE}/{doc}{SCOPE}").json()
    assert info['save_error'] and info['pending_save'] == request and info['revision'] == 2
    assert callback(client, doc, {'status': 4, 'key': key}).json() == {'error': 1}
    assert client.get(f"{BASE}/{doc}/download{SCOPE}").content == original
    assert client.post(f"{BASE}/{doc}/release-unopened{SCOPE}").json() == {'released': False}


def test_audit_is_transactional_immutable_and_paged(api):
    client, engine, users, current, monkey = api
    doc = upload(client)
    with Session(engine) as db:
        record = db.get(Document, doc)
        for i in range(55):
            management.audit(db, record, users['owner'], 'rename', {'from': str(i), 'to': str(i+1)})
        db.commit()
    first = client.get(f"{BASE}/{doc}/events{SCOPE}").json()
    second = client.get(f"{BASE}/{doc}/events{SCOPE}&before={first['next_before']}").json()
    assert len(first['items']) == 50 and len(second['items']) == 6
    assert not set(e['id'] for e in first['items']) & set(e['id'] for e in second['items'])
    for action in ('UPDATE online_document_events SET kind=\'x\'', 'DELETE FROM online_document_events'):
        with Session(engine) as db, pytest.raises(DatabaseError): db.execute(text(action))
    with engine.connect() as conn: management.verify_evidence_guards(conn)
    with engine.begin() as conn: conn.execute(text('DROP TRIGGER online_document_events_no_update'))
    with engine.connect() as conn, pytest.raises(RuntimeError): management.verify_evidence_guards(conn)


def test_restore_protected_history_recovers_policy_and_can_disable(api):
    client, engine, users, current, monkey = api
    doc = upload(client); enabled = enable(client, doc)
    client.put(f"{BASE}/{doc}/protection{SCOPE}", json=payload(client, doc, enabled=False))
    result = client.post(f"{BASE}/{doc}/restore{SCOPE}", json=payload(client, doc, revision=2))
    assert result.status_code == 200 and result.json()['protected'] is True
    assert client.get(f"{BASE}/{doc}/download{SCOPE}").content == enabled
    with Session(engine) as db:
        assert db.get(Document, doc).protection['baseline_revision'] == 4
    assert client.put(f"{BASE}/{doc}/protection{SCOPE}", json=payload(client, doc, enabled=False)).json()['protected'] is False


def test_idempotent_retry_after_later_change_reports_current_state_separately(api):
    client, engine, users, current, monkey = api
    doc = upload(client); request = payload(client, doc, revision=1)
    client.post(f"{BASE}/{doc}/restore{SCOPE}", json=request)
    enable(client, doc)
    retry = client.post(f"{BASE}/{doc}/restore{SCOPE}", json=request)
    assert retry.status_code == 200 and retry.json()['operation_notice']
    assert retry.json()['revision'] == 3


@pytest.mark.parametrize('ranges', [['A1:XFD1048576'], ['Z3:A1'], ['B0'], ['B1:B9999999']])
def test_protection_rejects_oversized_or_invalid_ranges_without_change(api, ranges):
    client, engine, users, current, monkey = api
    doc = upload(client)
    result = client.put(f"{BASE}/{doc}/protection{SCOPE}", json=payload(client, doc, enabled=True, rules=[{'sheet': '进出货', 'ranges': ranges}]))
    assert result.status_code == 422
    assert client.get(f"{BASE}/{doc}{SCOPE}").json()['revision'] == 1


def test_protection_rejects_input_cells_that_exceed_final_book_limit(api):
    client, engine, users, current, monkey = api
    doc = upload(client)
    original = client.get(f"{BASE}/{doc}/download{SCOPE}").content
    monkey.setattr(protection, 'MAX_CELLS', 4)
    result = client.put(f"{BASE}/{doc}/protection{SCOPE}", json=payload(client, doc, enabled=True, rules=[{'sheet': '进出货', 'ranges': ['D1']}]))
    assert result.status_code == 422
    assert client.get(f"{BASE}/{doc}{SCOPE}").json()['revision'] == 1
    assert client.get(f"{BASE}/{doc}/download{SCOPE}").content == original
