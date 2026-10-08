from io import BytesIO
from concurrent.futures import ThreadPoolExecutor

import pytest
from openpyxl import load_workbook

from test_molding_sample_api import make_client, login_as
from test_carton_master_api import prepare, read, payload
from test_carton_transaction_guards_api import BASE


def excel(kind, rows):
    from app.services.carton_master_import import template
    book = load_workbook(BytesIO(template(kind)))
    for row in rows:
        book["导入数据"].append(row)
    result = BytesIO()
    book.save(result)
    return result.getvalue()


def upload(client, kind, content, token=None, factory="huaxing", filename="模板.xlsx"):
    return client.post(BASE + f"/master-data/import/{kind}/{'apply' if token is not None else 'preview'}",
                       data={"factory_id": factory, "preview_token": token or ""},
                       files={"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def config(code="ITEM-001", group="标准", count=24, paper="外箱", preferred="否", status="启用", note=""):
    return [code, "产品", group, paper, "A33", 30, 20, 15, "cm", "个", count, preferred, status, note]


def test_templates_preview_only_and_options_preserve_rules_restore_hidden(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        for kind in ["paper-options", "configurations"]:
            response = client.get(BASE + f"/master-data/import/{kind}/template", params={"factory_id": "huaxing"})
            assert response.status_code == 200
            assert "filename*=UTF-8''" in response.headers["content-disposition"]
            book = load_workbook(BytesIO(response.content))
            assert book.sheetnames == ["导入数据", "说明与示例"]
            assert book["导入数据"].max_row == 1
            headers = [c.value for c in book["导入数据"][1]]
            assert all(label in headers for label in ["纸品类型", "纸质", "长", "宽", "高"])
            assert "规格" not in headers
            if kind == "configurations":
                examples = list(book["说明与示例"].iter_rows(values_only=True))[-2:]
                assert [row[0] for row in examples] == ["00123", "00123"]
                assert [row[2] for row in examples] == ["标准装", "标准装"]
                assert [row[3] for row in examples] == ["外箱", "内箱"]
                assert [row[13] for row in examples] == ["同一货号多纸品示例", "同一货号多纸品示例"]
            assert upload(client, kind, response.content).json()["errors"]
        original = client.post(BASE + "/master-data", json={"factory_id": "huaxing", "kind": "RULE", "code": "",
            "data": {"lead_days": 9, "paper_types": ["外箱"], "hidden_paper_types": ["底卡"], "note": "保留说明"},
            "status": "INACTIVE", "reason": "维护默认规则"}).json()
        content = excel("paper-options", [["底卡", "A33", 30.0, 20, 15], [" 外箱 ", "a33", None, None, None]])
        before = read(client)["records"]
        preview = upload(client, "paper-options", content).json()
        assert preview["errors"] == []
        assert (preview["added"], preview["skipped"]) == (3, 2)
        assert read(client)["records"] == before
        applied = upload(client, "paper-options", content, preview["preview_token"])
        assert applied.status_code == 200, applied.text
        row = read(client)["records"][0]
        assert row["data"]["lead_days"] == 9 and row["data"]["note"] == "保留说明"
        assert row["status"] == "INACTIVE" and row["revision"] == original["revision"] + 1
        assert row["data"]["hidden_paper_types"] == []
        assert row["data"]["specifications"] == ["30*20*15"]
        assert row["data"]["paper_qualities"] == ["A33"]
        assert upload(client, "paper-options", content, preview["preview_token"]).status_code == 409
        assert upload(client, "paper-options", content).json()["added"] == 0


def test_configuration_template_examples_apply_as_one_multi_paper_configuration(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        template_response = client.get(
            BASE + "/master-data/import/configurations/template",
            params={"factory_id": "huaxing"},
        )
        book = load_workbook(BytesIO(template_response.content))
        examples = list(book["说明与示例"].iter_rows(values_only=True))[-2:]
        for row in examples:
            book["导入数据"].append(row)
        content = BytesIO()
        book.save(content)

        preview = upload(client, "configurations", content.getvalue()).json()
        assert preview["errors"] == []
        assert (preview["added"], preview["skipped"]) == (1, 0)
        assert preview["details"] == ["00123 / 标准装：新增 2 行纸品"]

        applied = upload(client, "configurations", content.getvalue(), preview["preview_token"])
        assert applied.status_code == 200, applied.text
        records = read(client)["records"]
        assert len(records) == 1
        assert records[0]["code"] == "00123"
        assert len(records[0]["data"]["lines"]) == 2


def test_configurations_multi_paper_variants_duplicate_skip_and_audit(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        content = excel("configurations", [config(group="", preferred="是"), config(group="", paper="底卡", preferred="是"),
                                             config(group="大包装", count=48), config(group="大包装", count=48, paper="底卡")])
        preview = upload(client, "configurations", content).json()
        assert preview["errors"] == [] and preview["added"] == 2
        assert read(client)["records"] == []
        applied = upload(client, "configurations", content, preview["preview_token"])
        assert applied.status_code == 200, applied.text
        records = read(client)["records"]
        assert len(records) == 2
        assert {r["data"]["packing_name"] for r in records} == {"标准装", "大包装（48个装）"}
        assert all(r["customer_code"] == "" and len(r["data"]["lines"]) == 2 for r in records)
        assert sum(r["preferred"] for r in records) == 1
        changed = excel("configurations", [config(group="", status="停用", note="不会覆盖"), config(group="", paper="底卡", status="停用", note="不会覆盖")])
        duplicate = upload(client, "configurations", changed).json()
        assert duplicate["added"] == 0 and duplicate["skipped"] == 1
        assert upload(client, "configurations", changed, duplicate["preview_token"]).status_code == 200
        assert read(client)["records"] == records
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonAuditEvent
        from sqlalchemy import select
        with SessionLocal() as db:
            events = list(db.scalars(select(CartonAuditEvent).where(CartonAuditEvent.event_type == "MASTER_DATA_SAVED")))
            assert len(events) == 2


def test_rejects_changed_file_factory_kind_master_and_missing_confirmation(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        content = excel("paper-options", [["底卡", None, None, None, None]])
        preview = upload(client, "paper-options", content).json()
        token = preview["preview_token"]
        assert upload(client, "paper-options", content, "").status_code == 422
        assert upload(client, "paper-options", content, token, factory="huadeng").status_code == 409
        assert upload(client, "paper-options", excel("paper-options", [["内箱", None, None, None, None]]), token).status_code == 409
        assert upload(client, "configurations", excel("configurations", [config()]), token).status_code == 409
        client.post(BASE + "/master-data", json={"factory_id": "huaxing", "kind": "WORKSHOP", "code": "新增车间", "data": {}, "reason": "维护车间资料"})
        assert upload(client, "paper-options", content, token).status_code == 409
        assert not any(r["kind"] == "RULE" for r in read(client)["records"])


@pytest.mark.parametrize("row, expected", [
    (["外箱、底卡", None, None, None, None], "单一项"),
    (["外箱", "A33;B33", None, None, None], "单一项"),
    (["外箱", None, 30, None, 15], "全部填写"),
    (["外箱", None, 30, 0, 15], "正数"),
    (["外箱", None, "=10+20", 20, 15], "公式"),
    (["长" * 65, None, None, None, None], "64"),
])
def test_invalid_rows_never_partially_apply(monkeypatch, row, expected):
    with make_client(monkeypatch) as client:
        prepare(client)
        content = excel("paper-options", [["有效项", None, None, None, None], row])
        preview = upload(client, "paper-options", content).json()
        assert expected in "；".join(preview["errors"])
        assert upload(client, "paper-options", content, preview["preview_token"]).status_code == 422
        assert read(client)["records"] == []


def test_group_conflicts_and_bounds(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        for rows, expected in [([config(), config()], "重复"), ([config(), config(paper="底卡", note="冲突")], "冲突"),
                               ([config(group="甲"), config(group="乙")], "相同配置"),
                               ([config(preferred="是"), config(group="乙", count=48, preferred="是")], "推荐")]:
            result = upload(client, "configurations", excel("configurations", rows)).json()
            assert expected in "；".join(result["errors"])
        assert upload(client, "paper-options", b"bad").status_code == 422
        assert upload(client, "paper-options", b"x" * (5 * 1024 * 1024 + 1)).status_code == 422
        assert upload(client, "paper-options", excel("paper-options", []), filename="old.xls").status_code == 422
        assert upload(client, "paper-options", excel("paper-options", [["外箱"]] * 1001)).status_code == 422


def test_permissions_and_factory_isolation(monkeypatch):
    monkeypatch.setenv("AUTHZ_MODE", "enforce")
    with make_client(monkeypatch) as client:
        prepare(client)
        profile = login_as(client, "warehouse_keeper")
        assert read(client)["can_manage"]
        imports = [("paper-options", [["外箱"]]), ("configurations", [config()]), ("locations", [["A车间", "A1-A3"]])]
        for kind, rows in imports:
            content = excel(kind, rows)
            assert client.get(BASE + f"/master-data/import/{kind}/template", params={"factory_id": "huaxing"}).status_code == 200
            preview = upload(client, kind, content)
            assert preview.status_code == 200 and preview.json()["errors"] == []
            assert client.get(BASE + f"/master-data/import/{kind}/template", params={"factory_id": "huadeng"}).status_code == 403
            assert upload(client, kind, content, factory="huadeng").status_code == 403
            assert upload(client, kind, content, preview.json()["preview_token"], factory="huadeng").status_code == 403
            if kind == "locations":
                assert upload(client, kind, content, preview.json()["preview_token"]).status_code == 200
        # Exercise the real master permission gate independently of factory scope.
        from app.db import SessionLocal
        from app.models.auth import AuthPermission, AuthUserPermissionOverride
        from sqlalchemy import select
        with SessionLocal() as db:
            permission = db.scalar(select(AuthPermission).where(AuthPermission.code == "carton_procurement:master_manage"))
            db.add(AuthUserPermissionOverride(id="master-import-deny", user_id=profile["id"],
                permission_id=permission.id, effect="deny", factory_id="huaxing", department="*"))
            db.commit()
        assert not read(client)["can_manage"]
        for kind, rows in imports:
            content = excel(kind, rows)
            assert upload(client, kind, content).status_code == 403
            assert upload(client, kind, content, "token").status_code == 403
            assert client.get(BASE + f"/master-data/import/{kind}/template", params={"factory_id": "huaxing"}).status_code == 403


def test_item_numbers_require_real_text_and_preserve_leading_zeroes(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        from app.services.carton_master_import import template
        book = load_workbook(BytesIO(template("configurations")))
        example = book["说明与示例"].cell(book["说明与示例"].max_row - 1, 1)
        assert example.value == "00123" and example.data_type == "s"
        assert any("货号必须是真实文本" in str(row[0].value) for row in book["说明与示例"])
        text_row = config("00123")
        text_row[1] = 2026  # Product names retain their existing scalar handling.
        content = excel("configurations", [text_row])
        preview = upload(client, "configurations", content).json()
        assert preview["errors"] == []
        assert upload(client, "configurations", content, preview["preview_token"]).status_code == 200
        saved = read(client)["records"][0]
        assert saved["code"] == "00123" and saved["data"]["product_name"] == "2026"
        for value, number_format in [(123, "00000"), (12345678901234567890, "General")]:
            invalid = load_workbook(BytesIO(excel("configurations", [config("VALID"), config(value)])))
            invalid["导入数据"]["A3"].number_format = number_format
            output = BytesIO()
            invalid.save(output)
            preview = upload(client, "configurations", output.getvalue()).json()
            assert preview["errors"] == ["第 3 行：货号必须是真实文本，不能使用数字单元格或补零显示格式；请先设为文本后重新输入，或以单引号开头输入"]
            assert upload(client, "configurations", output.getvalue(), preview["preview_token"]).status_code == 422
            assert [r["code"] for r in read(client)["records"]] == ["00123"]


def test_save_failure_rolls_back_whole_batch(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        content = excel("configurations", [config("ITEM-1"), config("ITEM-2")])
        preview = upload(client, "configurations", content).json()
        from app.services import carton_master
        from fastapi import HTTPException
        original = carton_master.save_record
        calls = 0
        def fail_second(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise HTTPException(409, "模拟保存冲突")
            return original(*args, **kwargs)
        monkeypatch.setattr(carton_master, "save_record", fail_second)
        assert upload(client, "configurations", content, preview["preview_token"]).status_code == 409
        assert read(client)["records"] == []


def test_concurrent_confirmations_share_factory_lock(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        content = excel("configurations", [config()])
        preview = upload(client, "configurations", content).json()
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(lambda _: upload(client, "configurations", content, preview["preview_token"]), range(2)))
        assert sorted(r.status_code for r in responses) in ([200, 409], [200, 429])
        # The bounded file worker may reject before entering the factory lock.
        # Once it is free, the original stale preview must still be rejected.
        assert upload(client, "configurations", content, preview["preview_token"]).status_code == 409
        assert len(read(client)["records"]) == 1


def test_locations_template_range_preview_apply_preserves_existing(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        template = client.get(BASE + "/master-data/import/locations/template", params={"factory_id": "huaxing"})
        assert template.status_code == 200 and "filename*=UTF-8''" in template.headers["content-disposition"]
        book = load_workbook(BytesIO(template.content))
        assert [c.value for c in book["导入数据"][1]] == ["仓库", "仓位或范围"]
        assert book["导入数据"].max_row == 1
        assert upload(client, "locations", template.content).json()["errors"]
        first = client.post(BASE + "/inventory/locations", json={"factory_id": "huaxing", "warehouse": "a车间", "bin_code": "a1", "reason": "建立测试仓位"}).json()
        second = client.post(BASE + "/inventory/locations", json={"factory_id": "huaxing", "warehouse": "A车间", "bin_code": "B01", "reason": "建立测试仓位"}).json()
        stopped = client.patch(BASE + "/inventory/locations/" + second["id"], json={"factory_id": "huaxing", "warehouse": "A车间", "bin_code": "B01", "status": "INACTIVE", "expected_revision": second["revision"], "reason": "停用测试仓位"})
        assert stopped.status_code == 200
        from app.db import SessionLocal
        from app.models.carton_positions import CartonPositionEntry
        from app.models.carton_procurement import CartonAuditEvent
        from sqlalchemy import select
        import json
        from decimal import Decimal
        with SessionLocal() as db:
            db.add(CartonPositionEntry(factory_id="huaxing", inventory_key="test-existing-stock", location_id=first["id"], transfer_id="test-existing-transfer", quantity=Decimal("7"), occurred_at="2026-09-01"))
            db.commit()
        before = read(client)["locations"]
        content = excel("locations", [[" a车间 ", "a1-a25、b01-b03"], ["新仓", "001"]])
        preview = upload(client, "locations", content).json()
        assert preview["errors"] == [] and (preview["added"], preview["skipped"]) == (27, 2)
        assert read(client)["locations"] == before
        assert "停用" in "；".join(preview["details"])
        applied = upload(client, "locations", content, preview["preview_token"])
        assert applied.status_code == 200, applied.text
        after = read(client)["locations"]
        assert len(after) == len(before) + 27
        assert [r for r in after if r["id"] in {r["id"] for r in before}] == before
        assert any(r["warehouse"] == "新仓" and r["bin_code"] == "001" and r["status"] == "ACTIVE" for r in after)
        assert {r["bin_code"] for r in after if r["warehouse"] == "A车间"} == {*(f"A{i}" for i in range(1, 26)), "B01", "B02", "B03"}
        with SessionLocal() as db:
            entries = list(db.scalars(select(CartonPositionEntry)))
            assert len(entries) == 1 and entries[0].quantity == Decimal("7") and entries[0].location_id == first["id"]
            audits = [e for e in db.scalars(select(CartonAuditEvent).where(CartonAuditEvent.event_type == "INVENTORY_LOCATION_CREATED")) if json.loads(e.detail_json).get("reason") == "模板导入仓库仓位"]
            assert len(audits) == 27
            assert all(e.factory_id == "huaxing" and json.loads(e.detail_json)["status"] == "ACTIVE" for e in audits)
        duplicate = upload(client, "locations", content).json()
        assert (duplicate["added"], duplicate["skipped"]) == (0, 29)
        assert upload(client, "locations", content, duplicate["preview_token"]).status_code == 200
        assert read(client)["locations"] == after


def test_locations_invalid_expressions_and_expanded_limit_are_atomic(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        before = read(client)["locations"]
        invalid = [([[123, "A1"]], "真实文本"), ([["仓", 123]], "真实文本"), ([["仓", "=1+1"]], "公式"),
                   ([["仓", None]], "真实文本"), ([["仓、另一仓", "A1"]], "单一项"), ([["仓", "A3-A1"]], "逆序"),
                   ([["仓", "A1-B3"]], "同前缀"), ([["仓", "A1-A03"]], "宽度"), ([["仓", "A1,,A2"]], "空项"),
                   ([["仓", "A1-A1001"]], "1000"), ([["仓", "A1-A600"], ["仓", "B1-B401"]], "1000"),
                   ([["仓", "A1-A2、A2"]], "重复"), ([[" 仓 ", "a1"], ["仓", "A1"]], "重复"),
                   ([["待核仓位", "A1"]], "系统保留"), ([["仓", "CL-UNKNOWN-123"]], "系统保留"),
                   ([["CL-UNKNOWN-X", "A1"]], "系统保留"), ([["仓", "待核仓位"]], "系统保留")]
        for rows, expected in invalid:
            content = excel("locations", rows)
            preview = upload(client, "locations", content).json()
            assert expected in "；".join(preview["errors"]), (rows, preview)
            assert upload(client, "locations", content, preview["preview_token"]).status_code == 422
        assert read(client)["locations"] == before
        boundary = upload(client, "locations", excel("locations", [["仓", "A1-A1000"]])).json()
        assert boundary["errors"] == [] and boundary["added"] == 1000
        assert read(client)["locations"] == before  # Even a maximal preview remains read-only.
        legacy = upload(client, "locations", excel("locations", [["默认仓", "A-00"], ["新仓", "A-01-A-03"]])).json()
        assert legacy["errors"] == [] and legacy["skipped"] == 1 and legacy["added"] == 3
        assert all(f"A-{i:02}" in "；".join(legacy["details"]) for i in range(1, 4))


def test_locations_snapshot_changes_invalidate_preview_and_concurrent_apply(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        from app.db import SessionLocal
        from app.models.carton_positions import CartonLocation
        from app.services.carton_procurement import _lock_receipt_factory
        content = excel("locations", [["导入仓", "A1-A2"]])
        preview = upload(client, "locations", content).json()
        row = client.post(BASE + "/inventory/locations", json={"factory_id": "huaxing", "warehouse": "变更仓", "bin_code": "X1", "reason": "新增并发仓位"}).json()
        assert upload(client, "locations", content, preview["preview_token"]).status_code == 409
        # Isolate each snapshot field, including changes independent of revision.
        for field, value in [("warehouse", "改名仓"), ("bin_code", "X2"), ("status", "INACTIVE"), ("revision", 7), ("id", "CL-REPLACED-TEST")]:
            preview = upload(client, "locations", content).json()
            with SessionLocal() as db:
                _lock_receipt_factory(db, "huaxing")
                existing = db.get(CartonLocation, row["id"])
                setattr(existing, field, value)
                db.commit()
            assert upload(client, "locations", content, preview["preview_token"]).status_code == 409, field
            if field == "id":
                row["id"] = value
        preview = upload(client, "locations", content).json()
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(lambda _: upload(client, "locations", content, preview["preview_token"]), range(2)))
        assert sorted(r.status_code for r in responses) in ([200, 409], [200, 429])
        assert upload(client, "locations", content, preview["preview_token"]).status_code == 409
        assert len([r for r in read(client)["locations"] if r["warehouse"] == "导入仓"]) == 2


def test_locations_failure_rolls_back_rows_and_audit(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonAuditEvent
        from app.services import carton_positions
        from sqlalchemy import select
        from fastapi import HTTPException
        before = read(client)["locations"]
        with SessionLocal() as db:
            before_audit = list(db.scalars(select(CartonAuditEvent.id).where(CartonAuditEvent.event_type == "INVENTORY_LOCATION_CREATED").order_by(CartonAuditEvent.id)))
        content = excel("locations", [["新仓", "A1-A3"]])
        preview = upload(client, "locations", content).json()
        original = carton_positions.create_location
        calls = 0
        def fail_second(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise HTTPException(409, "模拟中途失败")
            return original(*args, **kwargs)
        monkeypatch.setattr(carton_positions, "create_location", fail_second)
        assert upload(client, "locations", content, preview["preview_token"]).status_code == 409
        assert read(client)["locations"] == before
        with SessionLocal() as db:
            assert list(db.scalars(select(CartonAuditEvent.id).where(CartonAuditEvent.event_type == "INVENTORY_LOCATION_CREATED").order_by(CartonAuditEvent.id))) == before_audit


def test_paper_weight_import_and_legacy_template_compatibility(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        content = excel("configurations", [config(paper="外箱") + ["8.125", "9.25"], config(paper="内箱") + ["0.4", "0.5"]])
        preview = upload(client, "configurations", content).json()
        assert preview["errors"] == []
        assert upload(client, "configurations", content, preview["preview_token"]).status_code == 200
        papers = read(client)["records"][0]["data"]["lines"]
        assert {paper["packaging_type"]: (paper["net_weight_kg"], paper["gross_weight_kg"]) for paper in papers} == {"外箱": ("8.125", "9.25"), "内箱": ("0.4", "0.5")}
        invalid = upload(client, "configurations", excel("configurations", [config(code="BAD") + ["5", "4"]])).json()
        assert invalid["errors"] and "毛重" in invalid["errors"][0]
        old = load_workbook(BytesIO(excel("configurations", [config(code="LEGACY")])))
        old["导入数据"].delete_cols(15, 2)
        result = BytesIO(); old.save(result)
        legacy = upload(client, "configurations", result.getvalue()).json()
        assert legacy["errors"] == [] and legacy["added"] == 1
        assert upload(client, "configurations", result.getvalue(), legacy["preview_token"]).status_code == 200
