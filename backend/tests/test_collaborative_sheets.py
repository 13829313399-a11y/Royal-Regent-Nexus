"""ACL, state/CAS, source immutability and exact workbook preservation regressions."""
from dataclasses import replace
from io import BytesIO
import importlib.util
from pathlib import Path
import struct
from zipfile import ZipFile
from unittest.mock import Mock

from alembic.migration import MigrationContext
from alembic.operations import Operations
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
import openpyxl
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import PatternFill, Font
from PIL import Image
import pytest
from sqlalchemy import create_engine, event as sql_event, inspect, select, text
from sqlalchemy.exc import DatabaseError
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
import xlrd
import xlwt

from app.api.collaborative_sheets import router
from app.core.config import settings
from app.db import get_db
from app.models.auth import AuthUser, EmployeeProfile
from app.models.identity import EmployeeAssignment, IamOrgUnit, IamOrgDepartment
from app.models.collaborative_sheets import CollaborativeSheet as Task, CollaborativeSheetEvent as Event, CollaborativeSheetSubmission as Submission
from app.schemas.collaborative_sheets import CellsInput, GrantsInput
from app.services.auth import AuthContext, AuthProfileContext, get_current_user
from app.services import collaborative_sheets as service
from app.services import collaborative_sheet_files as files

BASE = "/api/tools/collaborative-sheets"
SCOPE = "?factory_id=huaxing"
MODELS = (Task, Event, Submission)


def context(identity="owner", factory="huaxing", department="engineering"):
    return AuthContext(id=identity, username=identity, display_name=identity, roles=(), role_codes=(),
                       permissions=frozenset(), factory_scopes=(factory,), department_scopes=(department,),
                       profile=AuthProfileContext(primary_factory_id=factory, primary_department=department, confirmation_status="confirmed"))


def xlsx():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "填写区"
    sheet["A1"] = "原表说明"
    sheet["A2"], sheet["B2"], sheet["C2"] = "001", 12, "=B2*2"
    sheet["A2"].fill = PatternFill("solid", fgColor="00AACC")
    sheet["A2"].font = Font(name="宋体", size=14, bold=True)
    sheet.merge_cells("A4:B4")
    sheet["D6"] = "末格"
    sheet.row_dimensions[2].height = 27
    sheet.column_dimensions["B"].width = 20
    sheet.print_area = "A1:D6"
    picture = BytesIO()
    Image.new("RGB", (24, 16), "red").save(picture, "PNG")
    picture.seek(0)
    sheet.add_image(SheetImage(picture), "D2")
    book.create_sheet("保留页")["A1"] = "=SUM('填写区'!B2)"
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def xls():
    book = xlwt.Workbook()
    sheet = book.add_sheet("填写区")
    sheet.write(0, 0, "原表说明")
    sheet.write(1, 0, "001", xlwt.easyxf("font: bold on; pattern: pattern solid, fore_colour yellow;"))
    sheet.write(1, 1, 12)
    sheet.write(1, 2, xlwt.Formula("B2*2"))
    sheet.write_merge(3, 3, 0, 1, "合并格")
    sheet.write(5, 3, "末格")
    sheet.row(1).height = 540
    sheet.col(1).width = 20 * 256
    other = book.add_sheet("保留页")
    other.write(0, 0, xlwt.Formula("SUM('填写区'!B2)"))
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


@pytest.fixture
def api(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "document_tools_enabled", True)
    monkeypatch.setattr(settings, "document_tools_storage_dir", str(tmp_path / "files"))
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    for model in (AuthUser, EmployeeProfile, IamOrgUnit, IamOrgDepartment, EmployeeAssignment, *MODELS):
        model.__table__.create(engine)
    with Session(engine) as db:
        for who in [context(), context("filler"), context("outsider", department="qc"), context("other", factory="huakang-a")]:
            db.add(AuthUser(id=who.id, username=who.username, display_name=who.display_name, password_salt="x", password_hash="x", status="active"))
            db.add(EmployeeProfile(user_id=who.id, primary_factory_id=who.profile.primary_factory_id,
                                   primary_department=who.profile.primary_department, confirmation_status="confirmed"))
        db.commit()
    app = FastAPI()
    app.include_router(router)
    active = {"user": context()}
    def session():
        with Session(engine) as db:
            yield db
    app.dependency_overrides[get_db] = session
    app.dependency_overrides[get_current_user] = lambda: active["user"]
    with TestClient(app) as client:
        yield client, active, engine
    engine.dispose()


def create(api, kind="xlsx"):
    result = api[0].post(BASE, data={"factory_id": "huaxing", "title": "周报"},
                         files={"file": ("原表." + kind, xls() if kind == "xls" else xlsx())})
    assert result.status_code == 201, result.text
    return result.json()


def endpoint(task, action=""):
    return BASE + "/" + task["id"] + action + SCOPE


def publish(api, kind="xlsx", principal="user", target="filler", region="A2:D6"):
    task = create(api, kind)
    result = api[0].put(endpoint(task, "/grants"), json={"expected_revision": 1, "grants": [
        {"principal_type": principal, "principal_id": target, "sheet": 0, "range": region}]})
    assert result.status_code == 200, result.text
    result = api[0].post(endpoint(task, "/state"), json={"expected_revision": 2, "status": "open"})
    assert result.status_code == 200, result.text
    return result.json()


def save(client, task, address="A2", value="0007", revision=None):
    return client.patch(endpoint(task, "/cells"), json={"expected_revision": revision or task["revision"],
                         "changes": [{"sheet": 0, "address": address, "value": value}]})


def test_draft_is_private_and_requires_explicit_valid_ranges(api):
    client, active, _ = api
    task = create(api)
    assert client.post(endpoint(task, "/state"), json={"expected_revision": 1, "status": "open"}).status_code == 422
    for principal, target, region in [("user", "other", "A2:D6"), ("department", "bad", "A2:D6"), ("user", "filler", "A4:A4"), ("user", "filler", "A1:Z999")]:
        r = client.put(endpoint(task, "/grants"), json={"expected_revision": 1, "grants": [
            {"principal_type": principal, "principal_id": target, "sheet": 0, "range": region}]})
        assert r.status_code == 422
    active["user"] = context("filler")
    assert client.get(endpoint(task)).status_code == 404
    assert client.get(BASE + SCOPE).json() == {"items": []}
    assert client.get(endpoint(task, "/download")).status_code == 404


@pytest.mark.parametrize("kind", ["xls", "xlsx"])
def test_fill_submit_close_reopen_and_same_format_owner_download(api, kind):
    client, active, engine = api
    task = publish(api, kind)
    with Session(engine) as db:
        row = db.get(Task, task["id"])
        original = service.source_bytes(row)
    active["user"] = context("filler")
    assert client.get(endpoint(task)).status_code == 200
    assert client.get(endpoint(task, "/download")).status_code == 403
    for action, method, payload in [("/state", client.post, {"expected_revision": 3, "status": "closed"}),
                                    ("/grants", client.put, {"expected_revision": 3, "grants": []})]:
        assert method(endpoint(task, action), json=payload).status_code == 403
    assert save(client, task, "A1", "outside").status_code == 403
    assert save(client, task, "C2", "overwrite formula").status_code == 422
    assert save(client, task, "B4", "merge child").status_code == 422
    assert save(client, task, "A2", "=HYPERLINK(1)").status_code == 422
    r = save(client, task)
    assert r.status_code == 200, r.text
    task = r.json()
    assert save(client, task, revision=3).status_code == 409
    submitted = client.post(endpoint(task, "/submit"), json={"expected_revision": 4})
    assert submitted.status_code == 200, submitted.text
    assert submitted.json()["submissions"][0]["current"] is True
    task = save(client, task, "B2", 20).json()
    assert task["submissions"][0]["current"] is False
    active["user"] = context()
    result = client.get(endpoint(task, "/download"))
    assert result.status_code == 200
    assert result.headers["x-workbook-revision"] == "5"
    if kind == "xls":
        book = xlrd.open_workbook(file_contents=result.content, formatting_info=True)
        assert book.sheet_by_index(0).cell_value(1, 0) == "0007"
        assert book.sheet_by_index(0).cell_value(1, 1) == 20
    else:
        book = openpyxl.load_workbook(BytesIO(result.content))
        assert book.active["A2"].value == "0007"
        assert book.active["C2"].value == "=B2*2"
    closed = client.post(endpoint(task, "/state"), json={"expected_revision": 5, "status": "closed"}).json()
    active["user"] = context("filler")
    assert save(client, closed).status_code == 409
    assert client.post(endpoint(closed, "/submit"), json={"expected_revision": 6}).status_code == 409
    active["user"] = context()
    assert client.post(endpoint(task, "/state"), json={"expected_revision": 6, "status": "open"}).status_code == 200
    with Session(engine) as db:
        row = db.get(Task, task["id"])
        assert service.source_bytes(row) == original
        assert db.scalars(select(Event).where(Event.task_id == row.id)).all()


def test_department_membership_and_cross_factory_never_inferred_from_roles(api):
    client, active, _ = api
    task = publish(api, principal="department", target="engineering")
    active["user"] = context("filler")
    assert client.get(endpoint(task)).status_code == 200
    active["user"] = context("filler", department="qc")
    assert client.get(endpoint(task)).status_code == 404
    active["user"] = replace(context("other", factory="huakang-a"), permissions=frozenset({"*"}), factory_scopes=("*",))
    assert client.get(endpoint(task)).status_code == 403
    assert client.get(endpoint(task).replace("huaxing", "huakang-a")).status_code == 404
    assert client.get(BASE + "/recipients" + SCOPE).status_code == 403
    active["user"] = context()
    assert {u["id"] for u in client.get(BASE + "/recipients" + SCOPE).json()["users"]} == {"owner", "filler", "outsider"}


def test_grant_revocation_conflicts_with_preloaded_writer(api):
    client, active, engine = api
    task = publish(api)
    with Session(engine) as stale_db:
        stale = stale_db.get(Task, task["id"])
        assert stale.revision == 3
        response = client.put(endpoint(task, "/grants"), json={"expected_revision": 3, "grants": [
            {"principal_type": "user", "principal_id": "outsider", "sheet": 0, "range": "A2:D6"}]})
        assert response.status_code == 200
        with pytest.raises(HTTPException) as error:
            service.save_cells(stale_db, stale, context("filler"), CellsInput(expected_revision=3, changes=[{"sheet": 0, "address": "A2", "value": "stale"}]))
        assert error.value.status_code == 409
    active["user"] = context("filler")
    assert client.get(endpoint(task)).status_code == 404


def test_all_or_nothing_save_and_history_append_only(api):
    client, active, engine = api
    task = publish(api)
    active["user"] = context("filler")
    result = client.patch(endpoint(task, "/cells"), json={"expected_revision": 3, "changes": [
        {"sheet": 0, "address": "A2", "value": "valid"}, {"sheet": 0, "address": "A1", "value": "forbidden"}]})
    assert result.status_code == 403
    with Session(engine) as db:
        assert db.get(Task, task["id"]).overrides == {}
        with pytest.raises(DatabaseError):
            db.execute(text("DELETE FROM collaborative_sheet_events"))


def add_account(engine, identity, department="sales-business", factory="huaxing", status="active", confirmation="confirmed"):
    with Session(engine) as db:
        db.add(AuthUser(id=identity, username=identity, display_name=identity, password_salt="test", password_hash="test", status=status))
        db.add(EmployeeProfile(user_id=identity, primary_factory_id=factory, primary_department=department, confirmation_status=confirmation))
        db.commit()


def test_participant_roster_resolves_all_current_department_accounts_and_deduplicates(api):
    client, active, engine = api
    for identity, kwargs in [
        ("business-a", {}), ("business-b", {}),
        ("disabled", {"status": "disabled"}), ("unconfirmed", {"confirmation": "needs_review"}),
        ("external-business", {"factory": "huakang-a"}), ("non-business", {"department": "qc"}),
    ]:
        add_account(engine, identity, **kwargs)
    task = publish(api, principal="department", target="sales-business")
    # Repeated ranges and overlapping account/department grants select people,
    # not extra roster rows. Owner is not assigned merely by owning the task.
    result = client.put(endpoint(task, "/grants"), json={"expected_revision": 3, "grants": [
        {"principal_type": "department", "principal_id": "sales-business", "sheet": 0, "range": "A2:D6"},
        {"principal_type": "department", "principal_id": "sales-business", "sheet": 0, "range": "A2:B2"},
        {"principal_type": "user", "principal_id": "business-a", "sheet": 0, "range": "A2:B2"},
    ]})
    assert result.status_code == 200, result.text
    roster = result.json()["participants"]
    assert [p["user_id"] for p in roster] == ["business-a", "business-b"]
    assert all(p["department"] == "sales-business" and p["status"] == "not_started" for p in roster)
    assert all(p["last_saved_at"] is None and p["submitted_at"] is None for p in roster)
    active["user"] = context("business-a", department="sales-business")
    assert client.get(endpoint(task, "/participants")).json()["participants"] == roster


def test_participant_status_transitions_follow_saves_and_document_revision(api):
    client, active, _ = api
    task = publish(api, principal="department", target="engineering")
    assert {p["user_id"]: p["status"] for p in task["participants"]} == {"filler": "not_started", "owner": "not_started"}
    active["user"] = context("filler")
    task = save(client, task).json()
    filling = next(p for p in task["participants"] if p["user_id"] == "filler")
    assert filling["status"] == "in_progress" and filling["last_saved_at"] and filling["submitted_at"] is None
    task = client.post(endpoint(task, "/submit"), json={"expected_revision": task["revision"]}).json()
    done = next(p for p in task["participants"] if p["user_id"] == "filler")
    assert done["status"] == "completed" and done["submitted_at"]
    assert done["last_saved_at"] == filling["last_saved_at"]
    active["user"] = context()
    task = save(client, task, "B2", 24).json()
    roster = {p["user_id"]: p for p in task["participants"]}
    assert roster["filler"]["status"] == "needs_confirmation"
    assert roster["filler"]["submitted_at"] == done["submitted_at"]
    assert roster["owner"]["status"] == "in_progress"
    active["user"] = context("filler")
    task = client.post(endpoint(task, "/submit"), json={"expected_revision": task["revision"]}).json()
    assert next(p for p in task["participants"] if p["user_id"] == "filler")["status"] == "completed"


def test_roster_keeps_old_save_history_beyond_recent_activity_window(api):
    client, active, engine = api
    task = publish(api)
    active["user"] = context("filler")
    task = save(client, task).json()
    saved_at = task["participants"][0]["last_saved_at"]
    with Session(engine) as db:
        for i in range(60):
            db.add(Event(id=f"later-{i}", task_id=task["id"], actor_id="owner", actor_name="owner", action="grants_changed",
                         revision=task["revision"], detail={}, created_at=f"9999-01-01T00:00:{i:02d}+00:00"))
        db.commit()
    result = client.get(endpoint(task)).json()
    assert len(result["activity"]) == 50 and all(e["action"] != "cells_saved" for e in result["activity"])
    assert result["participants"][0]["status"] == "in_progress"
    assert result["participants"][0]["last_saved_at"] == saved_at


def test_roster_removes_ineligible_members_and_accounts_for_new_membership(api):
    client, _, engine = api
    for identity in ("member-a", "member-b", "member-c", "member-d"):
        add_account(engine, identity)
    task = publish(api, principal="department", target="sales-business")
    assert len(task["participants"]) == 4
    with Session(engine) as db:
        db.get(AuthUser, "member-a").status = "disabled"
        db.get(EmployeeProfile, "member-b").primary_department = "qc"
        db.get(EmployeeProfile, "member-c").confirmation_status = "needs_review"
        db.get(EmployeeProfile, "member-d").primary_factory_id = "huakang-a"
        db.commit()
    assert client.get(endpoint(task, "/participants")).json()["participants"] == []
    add_account(engine, "member-new")
    assert [p["user_id"] for p in client.get(endpoint(task, "/participants")).json()["participants"]] == ["member-new"]
    # An explicit account grant survives an in-factory department transfer,
    # unlike department membership, while current department/name are shown.
    task = client.put(endpoint(task, "/grants"), json={"expected_revision": task["revision"], "grants": [
        {"principal_type": "user", "principal_id": "member-b", "sheet": 0, "range": "A2:D6"}]}).json()
    assert [(p["user_id"], p["department"]) for p in task["participants"]] == [("member-b", "qc")]


def test_roster_uses_canonical_v2_primary_identity_instead_of_stale_profile(api):
    client, active, engine = api
    for identity in ("v2-local", "v2-transferred", "v2-future", "v2-expired", "v2-old-epoch", "left-account"):
        add_account(engine, identity, department="engineering")
    with Session(engine) as db:
        for name, factory in (("local-org", "huaxing"), ("other-org", "huakang-a")):
            db.add(IamOrgUnit(id=name, name=name, kind="factory", legacy_factory_id=factory))
            db.add(IamOrgDepartment(org_unit_id=name, department_code="sales-business"))
        db.flush()
        for identity in ("v2-local", "v2-transferred", "v2-future", "v2-expired", "v2-old-epoch"):
            profile = db.get(EmployeeProfile, identity)
            profile.identity_mode = "v2"
            db.add(EmployeeAssignment(id="assignment-" + identity, user_id=identity,
                org_unit_id="other-org" if identity == "v2-transferred" else "local-org", department_code="sales-business",
                official_position_title="test", is_primary=True,
                valid_from="2099-01-01T00:00:00.000000Z" if identity == "v2-future" else "2020-01-01T00:00:00.000000Z",
                valid_until="2021-01-01T00:00:00.000000Z" if identity == "v2-expired" else None,
                employment_epoch=0 if identity == "v2-old-epoch" else 1, created_by="test", confirmed_by="test",
                created_at="2020-01-01T00:00:00.000000Z", updated_at="2020-01-01T00:00:00.000000Z"))
        db.get(EmployeeProfile, "left-account").employment_status = "left"
        db.commit()
    result = client.get(BASE + "/recipients" + SCOPE)
    assert result.status_code == 200, result.text
    users = {u["id"]: u for u in result.json()["users"]}
    assert users["v2-local"]["department"] == "sales-business"
    assert not {"v2-transferred", "v2-future", "v2-expired", "v2-old-epoch", "left-account"} & users.keys()
    task = publish(api, principal="department", target="sales-business")
    assert [p["user_id"] for p in task["participants"]] == ["v2-local"]
    with Session(engine) as db:
        db.get(IamOrgDepartment, ("local-org", "sales-business")).status = "inactive"
        db.commit()
    assert client.get(endpoint(task, "/participants")).json()["participants"] == []
    active["user"] = replace(context(), account_available=False)
    assert client.get(endpoint(task, "/participants")).status_code == 403


def test_roster_poll_is_lightweight_read_only_and_uses_existing_acl(api):
    client, active, engine = api
    task = create(api)
    active["user"] = context("filler")
    assert client.get(endpoint(task, "/participants")).status_code == 404


    active["user"] = context()
    task = publish(api)
    statements = []
    def capture(_connection, _cursor, statement, _params, _context, _many):
        statements.append(statement.lower())
    sql_event.listen(engine, "before_cursor_execute", capture)
    try:
        response = client.get(endpoint(task, "/participants"))
    finally:
        sql_event.remove(engine, "before_cursor_execute", capture)
    assert response.status_code == 200
    assert response.headers["cache-control"] == "private, no-store"
    assert set(response.json()) == {"id", "revision", "status", "participants"}
    assert response.json()["revision"] == task["revision"]
    assert not any(".manifest" in s or ".overrides" in s or ".detail" in s or ".password_hash" in s for s in statements)
    assert not any(s.lstrip().startswith(("insert", "update", "delete")) for s in statements)
    active["user"] = context("outsider", department="qc")
    assert client.get(endpoint(task, "/participants")).status_code == 404
    active["user"] = context("other", factory="huakang-a")
    assert client.get(endpoint(task, "/participants")).status_code == 403
    active["user"] = context()
    changed = client.put(endpoint(task, "/grants"), json={"expected_revision": task["revision"], "grants": [
        {"principal_type": "user", "principal_id": "outsider", "sheet": 0, "range": "A2:D6"}]}).json()
    assert [p["user_id"] for p in changed["participants"]] == ["outsider"]
    active["user"] = context("filler")
    assert client.get(endpoint(task, "/participants")).status_code == 404


def test_xlsx_exact_unedited_members_and_cell_style_formula_images():
    original = xlsx()
    manifest = files.inspect_workbook(original, "xlsx")
    assert manifest["sheets"][0]["images"][0]["url"].startswith("data:image/jpeg;")
    output = files.export_workbook(original, "xlsx", {"0:A2": "0009", "0:B3": 13.5, "0:A4": "合并填写"})
    with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(output)) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name not in {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}:
                assert before.read(name) == after.read(name), name
    before = openpyxl.load_workbook(BytesIO(original))
    after = openpyxl.load_workbook(BytesIO(output))
    assert after.active["A2"]._style == before.active["A2"]._style
    assert after.active["C2"].value == before.active["C2"].value
    assert after.active.merged_cells == before.active.merged_cells
    assert str(after.active.print_area) == str(before.active.print_area)
    assert after.active.row_dimensions[2].height == 27
    assert after.active.column_dimensions["B"].width == 20
    assert after.active["B3"].value == 13.5
    assert len(after.active._images) == 1
    assert files.export_workbook(original, "xlsx", {}) == original


def test_xls_formula_record_and_style_preserved_new_and_existing_cells():
    original = xls()
    output = files.export_workbook(original, "xls", {"0:A2": "0009", "0:B3": 13.5, "0:D6": None})
    a = xlrd.open_workbook(file_contents=original, formatting_info=True)
    b = xlrd.open_workbook(file_contents=output, formatting_info=True)
    assert b.sheet_by_index(0).cell_value(1, 0) == "0009"
    assert b.sheet_by_index(0).cell_value(2, 1) == 13.5
    assert b.sheet_by_index(0).cell_value(5, 3) == ""
    assert a.sheet_by_index(0).cell_xf_index(1, 0) == b.sheet_by_index(0).cell_xf_index(1, 0)
    assert a.sheet_by_index(0).merged_cells == b.sheet_by_index(0).merged_cells
    before_o, _, before, _ = files._xls_source(original)
    after_o, _, after, _ = files._xls_source(output)
    assert [v for _, t, v in files.records(before) if t == 6] == [v for _, t, v in files.records(after) if t == 6]
    assert [v for _, t, v in files.records(before) if t in {0xE0, 0x31, 0x7D}] == [v for _, t, v in files.records(after) if t in {0xE0, 0x31, 0x7D}]
    # Every generated Index points at a DBCell (xlwt commonly omits Index).
    for _, kind, value in files.records(after):
        if kind == 0x20B:
            for offset in range(16, len(value), 4):
                assert struct.unpack_from("<H", after, struct.unpack_from("<I", value, offset)[0])[0] == 0xD7
    before_o.close()
    after_o.close()


@pytest.mark.parametrize("kind", ["xls", "xlsx"])
def test_engine_refuses_formula_overwrite(kind):
    with pytest.raises(files.WorkbookError, match="公式"):
        files.export_workbook(xls() if kind == "xls" else xlsx(), kind, {"0:C2": "bad"})


def test_array_formula_result_cells_are_locked_in_manifest():
    from openpyxl.worksheet.formula import ArrayFormula
    book = openpyxl.Workbook()
    book.active["A1"] = 2
    book.active["A2"] = 3
    book.active["B1"] = ArrayFormula(ref="B1:B2", text="=A1:A2*2")
    book.active["B2"] = 0
    stream = BytesIO()
    book.save(stream)
    manifest = files.inspect_workbook(stream.getvalue(), "xlsx")
    cells = {c["address"]: c for c in manifest["sheets"][0]["cells"]}
    assert cells["B1"]["formula"] and cells["B2"]["formula"]


def test_integrity_check_refuses_changed_source(api):
    task = publish(api)
    with Session(api[2]) as db:
        row = db.get(Task, task["id"])
        from app.services.document_tools import storage
        storage.resolve(row.storage_key).write_bytes(b"tampered")
    assert api[0].get(endpoint(task, "/download")).status_code == 409


def rewrite_xlsx_member(original, path, transform):
    out = BytesIO()
    with ZipFile(BytesIO(original)) as src, ZipFile(out, "w") as dest:
        for item in src.infolist():
            value = src.read(item.filename)
            dest.writestr(item, transform(value) if item.filename == path else value)
    return out.getvalue()


def xlsx_full_column_format(original=None, minimum=1, maximum=16384):
    def mutate(data):
        root = files.xml(data)
        columns = root.find(f"{{{files.MAIN}}}cols")
        if columns is None:
            columns = files.ET.Element(f"{{{files.MAIN}}}cols")
            table = root.find(f"{{{files.MAIN}}}sheetData")
            root.insert(root.index(table), columns)
        columns.clear()
        files.ET.SubElement(columns, f"{{{files.MAIN}}}col", min=str(minimum), max=str(maximum),
                            width="19", customWidth="1", style="1", hidden="1")
        return files.ET.tostring(root)
    return rewrite_xlsx_member(original or xlsx(), "xl/worksheets/sheet1.xml", mutate)


@pytest.mark.parametrize("minimum", [1, 4, 27])
def test_xlsx_full_column_format_does_not_expand_grid_or_change_export(minimum):
    original = xlsx_full_column_format(minimum=minimum)
    manifest = files.inspect_workbook(original, "xlsx")
    sheet = manifest["sheets"][0]
    assert (sheet["rows"], sheet["columns"]) == (6, 4)
    assert set(sheet["column_widths"]) == {str(c - 1) for c in range(minimum, 5)}
    assert all(width == 19 * 7 + 5 for width in sheet["column_widths"].values())
    assert files.export_workbook(original, "xlsx", {}) == original
    output = files.export_workbook(original, "xlsx", {"0:A2": "updated"})
    with ZipFile(BytesIO(original)) as before, ZipFile(BytesIO(output)) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name not in {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}:
                assert before.read(name) == after.read(name), name
        cols = lambda data: files.ET.tostring(files.xml(data).find(f"{{{files.MAIN}}}cols"))
        assert cols(before.read("xl/worksheets/sheet1.xml")) == cols(after.read("xl/worksheets/sheet1.xml"))
    reopened = openpyxl.load_workbook(BytesIO(output))
    assert reopened.active["A2"].value == "updated"
    assert reopened.active["C2"].value == "=B2*2"
    assert reopened.active.max_column == 4
    assert reopened.active.column_dimensions[openpyxl.utils.get_column_letter(minimum)].max == 16384
    reopened.close()


@pytest.mark.parametrize("minimum,maximum", [(0, 16384), (5, 4), (1, 16385)])
def test_xlsx_invalid_column_metadata_still_rejected_before_loader(monkeypatch, minimum, maximum):
    original = xlsx_full_column_format(minimum=minimum, maximum=maximum)
    loader = Mock(side_effect=AssertionError("native loader must not run"))
    monkeypatch.setattr(files.openpyxl, "load_workbook", loader)
    with pytest.raises(files.WorkbookError):
        files.inspect_workbook(original, "xlsx")
    loader.assert_not_called()


@pytest.mark.parametrize("kind", ["cell", "merge", "hyperlink"])
def test_full_column_metadata_never_relaxes_allocating_range_limits(monkeypatch, kind):
    def mutate(data):
        root = files.xml(data)
        if kind == "cell":
            root.find(f".//{{{files.MAIN}}}c").set("r", "XFD1")
        else:
            parent, child = ("mergeCells", "mergeCell") if kind == "merge" else ("hyperlinks", "hyperlink")
            container = files.ET.SubElement(root, f"{{{files.MAIN}}}{parent}")
            files.ET.SubElement(container, f"{{{files.MAIN}}}{child}", ref="A1:XFD1048576")
        return files.ET.tostring(root)
    hostile = rewrite_xlsx_member(xlsx_full_column_format(), "xl/worksheets/sheet1.xml", mutate)
    loader = Mock(side_effect=AssertionError("native loader must not run"))
    monkeypatch.setattr(files.openpyxl, "load_workbook", loader)
    with pytest.raises(files.WorkbookError):
        files.inspect_workbook(hostile, "xlsx")
    loader.assert_not_called()


def test_xls_full_column_format_stays_bounded_and_preserved():
    book = xlwt.Workbook()
    sheet = book.add_sheet("Only")
    sheet.write(0, 0, "source")
    sheet.write(1, 2, xlwt.Formula("1+2"))
    for column in range(256):
        sheet.col(column).width = 19 * 256
    stream = BytesIO()
    book.save(stream)
    original = stream.getvalue()
    manifest = files.inspect_workbook(original, "xls")
    assert (manifest["sheets"][0]["rows"], manifest["sheets"][0]["columns"]) == (2, 3)
    assert set(manifest["sheets"][0]["column_widths"]) == {"0", "1", "2"}
    output = files.export_workbook(original, "xls", {"0:A1": "updated"})
    before_o, _, before, _ = files._xls_source(original)
    after_o, _, after, _ = files._xls_source(output)
    assert [v for _, t, v in files.records(before) if t in {0x7D, 0x6}] == [v for _, t, v in files.records(after) if t in {0x7D, 0x6}]
    before_o.close()
    after_o.close()
    assert xlrd.open_workbook(file_contents=output).sheet_by_index(0).cell_value(0, 0) == "updated"


@pytest.mark.parametrize("kind,refs", [
    ("merge", ["A1:XFD1048576"]), ("hyperlink", ["A1:XFD1048576"]),
    ("merge", ["B2:A1"]), ("hyperlink", ["A0:B2"]),
    ("merge", ["A1:D4", "B2:C3"]), ("hyperlink", ["A1:D4", "A1:D4"]),
    ("merge", ["A1:CV1001"]), ("hyperlink", ["A1:CV1001"]),
    ("array", ["A1:XFD1048576"]), ("dimension", ["A1:XFD1048576"]),
    ("cell", ["XFD1048576"]), ("bad_namespace_merge", ["A1:XFD1048576"]),
])
def test_xlsx_preflight_blocks_range_expansion_before_loader(monkeypatch, kind, refs):
    def mutate(data):
        root = files.xml(data)
        tag = lambda name: f"{{{files.MAIN}}}{name}"
        if kind in {"merge", "hyperlink", "bad_namespace_merge"}:
            container_name, name = ("hyperlinks", "hyperlink") if kind == "hyperlink" else ("mergeCells", "mergeCell")
            previous = root.find(tag(container_name))
            if previous is not None:
                root.remove(previous)
            container = files.ET.SubElement(root, tag(container_name))
            for ref in refs:
                files.ET.SubElement(container, "mergeCell" if kind == "bad_namespace_merge" else tag(name), ref=ref, **({"xmlns": ""} if kind == "bad_namespace_merge" else {}))
        elif kind == "array":
            cell = root.find(".//" + tag("c"))
            files.ET.SubElement(cell, tag("f"), t="array", ref=refs[0]).text = "A1"
        elif kind == "dimension":
            root.find(tag("dimension")).set("ref", refs[0])
        else:
            root.find(".//" + tag("c")).set("r", refs[0])
        return files.ET.tostring(root)
    hostile = rewrite_xlsx_member(xlsx(), "xl/worksheets/sheet1.xml", mutate)
    loader = Mock(side_effect=AssertionError("native loader must not run"))
    monkeypatch.setattr(files.openpyxl, "load_workbook", loader)
    with pytest.raises(files.WorkbookError):
        files.inspect_workbook(hostile, "xlsx")
    loader.assert_not_called()


def test_xlsx_preflight_counts_cumulative_effective_grid_without_dimension(monkeypatch):
    book = openpyxl.Workbook()
    book.active["A1"] = 1
    book.create_sheet("Second")["A1"] = 1
    stream = BytesIO()
    book.save(stream)
    def mutate(data):
        root = files.xml(data)
        root.remove(root.find(f"{{{files.MAIN}}}dimension"))
        merges = files.ET.SubElement(root, f"{{{files.MAIN}}}mergeCells")
        files.ET.SubElement(merges, f"{{{files.MAIN}}}mergeCell", ref="A1:CV600")
        return files.ET.tostring(root)
    hostile = stream.getvalue()
    for path in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
        hostile = rewrite_xlsx_member(hostile, path, mutate)
    loader = Mock(side_effect=AssertionError("native loader must not run"))
    monkeypatch.setattr(files.openpyxl, "load_workbook", loader)
    with pytest.raises(files.WorkbookError):
        files.inspect_workbook(hostile, "xlsx")
    loader.assert_not_called()


def test_xlsx_comment_range_is_rejected_before_loader(monkeypatch):
    from openpyxl.comments import Comment
    book = openpyxl.Workbook()
    book.active["A1"].comment = Comment("Note", "Author")
    stream = BytesIO()
    book.save(stream)
    with ZipFile(BytesIO(stream.getvalue())) as z:
        path = next(n for n in z.namelist() if "/comment" in n and n.endswith(".xml"))
    def mutate(data):
        root = files.xml(data)
        root.find(f"{{{files.MAIN}}}commentList/{{{files.MAIN}}}comment").set("ref", "A1:XFD1048576")
        return files.ET.tostring(root)
    hostile = rewrite_xlsx_member(stream.getvalue(), path, mutate)
    loader = Mock(side_effect=AssertionError("native loader must not run"))
    monkeypatch.setattr(files.openpyxl, "load_workbook", loader)
    with pytest.raises(files.WorkbookError):
        files.inspect_workbook(hostile, "xlsx")
    loader.assert_not_called()


def test_xlsx_ambiguous_relationship_cannot_bypass_preflight(monkeypatch):
    from copy import deepcopy
    def mutate(data):
        root = files.xml(data)
        root.append(deepcopy(root[0]))
        return files.ET.tostring(root)
    hostile = rewrite_xlsx_member(xlsx(), "xl/_rels/workbook.xml.rels", mutate)
    loader = Mock(side_effect=AssertionError("native loader must not run"))
    monkeypatch.setattr(files.openpyxl, "load_workbook", loader)
    with pytest.raises(files.WorkbookError):
        files.inspect_workbook(hostile, "xlsx")
    loader.assert_not_called()


@pytest.mark.parametrize("kind,ranges", [
    ("merge", [(0, 65535, 0, 255)]), ("hyperlink", [(0, 65535, 0, 255)]),
    ("merge", [(3, 1, 0, 2)]), ("merge", [(0, 999, 0, 100)]),
    ("merge", [(0, 4, 0, 4), (1, 2, 1, 2)]),
    ("hyperlink", [(0, 4, 0, 4), (0, 4, 0, 4)]),
    ("cell", [(65535, 65535, 255, 255)]),
    ("rich_text", [(65535, 65535, 255, 255)]),
])
def test_xls_preflight_blocks_implicit_ranges_before_loader(monkeypatch, kind, ranges):
    book = xlwt.Workbook()
    book.add_sheet("Only").write(0, 0, "seed")
    stream = BytesIO()
    book.save(stream)
    original = stream.getvalue()
    o, name, raw, bounds = files._xls_source(original)
    eof = next(p + bounds[0][1] for p, t, _ in files.records(raw[bounds[0][1]:]) if t == 0xA)
    if kind == "merge":
        inserted = files.rec(0xE5, struct.pack("<H", len(ranges)) + b"".join(struct.pack("<4H", *r) for r in ranges))
    elif kind == "hyperlink":
        inserted = b"".join(files.rec(0x1B8, struct.pack("<4H", *r) + bytes(24)) for r in ranges)
    elif kind == "rich_text":
        inserted = files.rec(0xD6, struct.pack("<3H", ranges[0][0], ranges[0][2], 0) + struct.pack("<HBH", 0, 0, 0))
    else:
        inserted = files.rec(0x201, struct.pack("<3H", ranges[0][0], ranges[0][2], 0))
    hostile = files._replace_ole_stream(original, o, name, raw[:eof] + inserted + raw[eof:])
    o.close()
    loader = Mock(side_effect=AssertionError("native loader must not run"))
    monkeypatch.setattr(files.xlrd, "open_workbook", loader)
    with pytest.raises(files.WorkbookError):
        files.inspect_workbook(hostile, "xls")
    loader.assert_not_called()


@pytest.mark.parametrize("value,mask,expected", [
    (-0.1358, "0.00%", "-13.58%"), (0.125, "0%", "13%"),
    (2426546.1800000006, "#,##0.00", "2,426,546.18"),
    (2426546.1800000006, "General", "2426546.18"),
    (-42.6, "#,##0.00;[Red](#,##0.00)", "(42.60)"),
    (12.5, "0.0#", "12.5"), (7, "0000", "0007"),
    ("0007", "0.00", "0007"), (True, "0.00%", "TRUE"),
    (2, "yyyy-mm-dd", "1900-01-02"), (0.5, "hh:mm:ss", "12:00:00"),
])
def test_simple_number_display_preserves_raw_value(value, mask, expected):
    assert files._display(value, mask) == expected


@pytest.mark.parametrize("kind", ["xls", "xlsx"])
def test_number_formats_in_manifest_and_overlay_keep_raw_values_and_source(kind):
    if kind == "xls":
        book = xlwt.Workbook()
        sheet = book.add_sheet("Fill")
        sheet.write(0, 0, -0.1358, xlwt.easyxf(num_format_str="0.00%"))
        sheet.write(0, 1, 2426546.1800000006, xlwt.easyxf(num_format_str="#,##0.00"))
    else:
        book = openpyxl.Workbook()
        book.active["A1"], book.active["B1"] = -0.1358, 2426546.1800000006
        book.active["A1"].number_format, book.active["B1"].number_format = "0.00%", "#,##0.00"
    out = BytesIO()
    book.save(out)
    original = out.getvalue()
    manifest = files.inspect_workbook(original, kind)
    cells = {c["address"]: c for c in manifest["sheets"][0]["cells"]}
    assert cells["A1"]["display"] == "-13.58%" and cells["A1"]["value"] == -0.1358
    assert cells["B1"]["display"] == "2,426,546.18"
    assert files._display(0.27, cells["A1"]["number_format"]) == "27.00%"
    assert files.export_workbook(original, kind, {}) == original


def test_startup_guard_and_additive_migration_preserve_existing_data(monkeypatch):
    from app import db as database
    path = Path(__file__).parents[1] / "alembic/versions/20261009_0121_collaborative_sheets.py"
    spec = importlib.util.spec_from_file_location("collaborative_migration", path)
    migration = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(migration)
    assert migration.down_revision == "20261005_0120"
    engine = create_engine("sqlite://", poolclass=StaticPool)
    monkeypatch.setattr(database, "engine", engine)
    database.ensure_collaborative_sheets_schema_ready()
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE existing_business (id INTEGER PRIMARY KEY, value TEXT)"))
        connection.execute(text("INSERT INTO existing_business VALUES (1, 'keep')"))
    with pytest.raises(RuntimeError, match="20261009_0121"):
        database.ensure_collaborative_sheets_schema_ready()
    with engine.begin() as connection:
        with Operations.context(MigrationContext.configure(connection)):
            migration.upgrade()
        assert connection.execute(text("SELECT value FROM existing_business")).scalar() == "keep"
    database.ensure_collaborative_sheets_schema_ready()
    with engine.begin() as connection:
        connection.execute(text("ALTER TABLE collaborative_sheets RENAME COLUMN grants TO old_grants"))
    with pytest.raises(RuntimeError, match="collaborative_sheets.grants"):
        database.ensure_collaborative_sheets_schema_ready()
