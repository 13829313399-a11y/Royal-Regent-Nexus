"""Warehouse totals omit PO work; review remains versioned, scoped and atomic."""
from io import BytesIO
from pathlib import Path
import openpyxl
import pytest
from fastapi import HTTPException
from sqlalchemy.orm import Session
from test_indonesia_materials import client, command, material, product, order, supply, historical_allocation, allocation, workspace, BASE
from test_indonesia_receipt_template import catalog, download
from test_indonesia_warehouse_import import choose, confirm, source
from app.models.indonesia_materials import IndonesiaMaterialRecord as Record
from app.services.indonesia_warehouse_import import parse_warehouse_receipts
from app.services.indonesia_receipt_template import SUMMARY_TITLE, HEADER_ROW, FIRST_ROW, SHEET


def summary(client, *materials, **changes):
    return download(client, mode='summary', selections=[], material_ids=[m['id'] for m in materials], **changes)


def sheet_rows(content):
    sheet = openpyxl.load_workbook(BytesIO(content))[SHEET]
    headers = [cell.value for cell in sheet[HEADER_ROW]]
    return sheet, [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=FIRST_ROW, values_only=True)]


def test_summary_aggregates_only_outstanding_same_material_and_keeps_shared_codes_separate(client):
    m, first = catalog(client)
    second = order(client, contract='TEST-B', item_no='00040149', quantity='200')
    order(client, contract='RRM', item_no='00040149', recipient='RRM', quantity='999')
    received = command(client, 'receipt_confirm', dict(actual_date='2026-10-22'), supply(client, m))
    allocation(client, received, first, normal='20', spare='5')
    other = material(client, code=m['data']['code'], name='同编码不同中文名')
    product(client, other, code='OTHER'); order(client, contract='OTHER', item_no='OTHER', quantity='50')
    content = summary(client, m, other).content
    sheet, rows = sheet_rows(content)
    assert sheet['A1'].value == SUMMARY_TITLE
    assert len(rows) == 2 and rows[0]['物料需求合计'] == '300'
    assert rows[0]['已实收并分配数量'] == '25' and rows[0]['剩余未送货数量'] == '275'
    assert rows[1]['剩余未送货数量'] == '50'
    assert not set(rows[0]) & {'合同号', '客户PO', '货号', '交付批次'}
    assert all(rows[0][h] is None for h in ['正常实收数量', '备品数量', '收料单号', '箱号'])
    assert sheet['I6'].protection.locked and not sheet['P6'].protection.locked
    assert sheet['I6'].fill.fgColor.rgb.endswith('F0F3F6') and sheet['P6'].fill.fgColor.rgb.endswith('FFF4CC')
    assert len(sheet_rows(summary(client, m, item_no='00040149').content)[1]) == 1
    summary(client, other, item_no='00040149', status=409)


def test_summary_clamps_each_demand_before_aggregation_and_excludes_expected_allocations(client):
    m, first = catalog(client)
    second = order(client, contract='TEST-B', item_no='00040149')
    first = command(client, 'demand_adjust', dict(demand_id=first['data']['demands'][0]['id'], quantity='75', reason='合成历史超覆盖'), first)
    received = command(client, 'receipt_confirm', dict(actual_date='2026-10-22'), supply(client, m))
    historical_allocation(client, received, first)
    expected = supply(client, m, batch='EXPECTED', purchase_no='EXPECTED')
    historical_allocation(client, expected, second)
    _, rows = sheet_rows(summary(client, m).content)
    assert rows[0]['物料需求合计'] == rows[0]['剩余未送货数量'] == '100'
    assert rows[0]['已实收并分配数量'] == '0'


@pytest.mark.parametrize('changes,status', [({'customer': '另一客人'}, 409), ({'recipient': 'RRM'}, 409), ({'factory_id': 'huadeng'}, 403), ({'material_ids': []}, 422)])
def test_summary_scope_and_empty_selection(client, changes, status):
    m, _ = catalog(client)
    body = dict(factory_id='huaxing', customer='合成客户', recipient='RRI', mode='summary', material_ids=[m['id']]) | changes
    response = client.post(BASE + '/warehouse-imports/template', json=body)
    assert response.status_code == status, response.text


def test_summary_return_priority_split_pack_confirm_and_export(client):
    m, first = catalog(client)
    second = order(client, contract='EARLIEST', item_no='00040149', delivery_date='2026-12-15')
    third = order(client, contract='FALLBACK', item_no='00040149', delivery_date='2026-11-02')
    with Session(client.engine) as db:
        for obj, required in [(first, '2026-11-03'), (second, '2026-11-01'), (third, None)]:
            row = db.get(Record, obj['id'])
            row.data = row.data | {'plan': (row.data['plan'] | {'required_date': required}) if required else None}
        db.commit()
    content = summary(client, m).content
    with pytest.raises(HTTPException, match='尚未填写本次实收数量'):
        parse_warehouse_receipts(content, '空模板.xlsx')
    book = openpyxl.load_workbook(BytesIO(content)); sheet = book[SHEET]
    headers = [cell.value for cell in sheet[HEADER_ROW]]
    for key, value in {'收料单号': 'WH-SUMMARY', '行号': '1', '实收日期': '2026-10-22', '正常实收数量': '110', '备品数量': '10'}.items():
        sheet.cell(FIRST_ROW, headers.index(key) + 1, value)
    stream = BytesIO(); book.save(stream)
    response = client.post(BASE + '/warehouse-imports/preview', data=dict(factory_id='huaxing', customer='合成客户', recipient='RRI', warehouse=''), files={'file': ('回传.xlsx', stream.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code == 200, response.text
    imported = response.json()['record']; row = imported['data']['rows'][0]
    assert not row['blocked'] and row['auto_material_id'] == m['id']
    assert not row['source'].get('po_no') and row['source']['normal_quantity'] == '110'
    targets = row['targets'][m['id']]
    assert [t['order_id'] for t in targets] == [second['id'], third['id'], first['id']]
    purchase = {key: source()[key] for key in ('purchase_no', 'purchase_line', 'purchase_date', 'supplier', 'price', 'currency')} | {'purchase_quantity': '120'}
    plans = [{k: t[k] for k in ('order_id', 'demand_id', 'order_version')} | dict(normal_quantity=n, spare_quantity=s) for t,n,s in [(targets[0], '100', '0'), (targets[1], '10', '10')]]
    posted = confirm(client, imported, [choose(imported, purchase=purchase, targets=plans)])
    data = workspace(client); shipment = next(r for r in data['records'] if r['id'] == posted['data']['shipment_ids'][0])
    payload = {k: shipment['data'][k] for k in ('customer', 'recipient', 'shipment_date', 'lines', 'packages', 'remarks')}
    assert len(payload['lines']) == 2
    payload['packages'] = [dict(label='汇总回料合箱', cartons=2, length_cm='30', width_cm='20', height_cm='10', carton_weight_kg='.02', items=payload['lines'])]
    packed = command(client, 'shipment_save', payload, shipment)
    done = command(client, 'shipment_confirm', {}, packed)
    export = client.get(BASE + f"/shipments/{done['id']}/export", params={'factory_id': 'huaxing'})
    assert export.status_code == 200 and done['data']['status'] == 'confirmed'
    assert sheet_rows(summary(client, m).content)[1][0]['剩余未送货数量'] == '180'
    out = Path('../.local-dev/receipt-summary'); out.mkdir(exist_ok=True, parents=True)
    (out/'汇总回料测试_确认出货.xlsx').write_bytes(export.content)
    (out/'汇总未收物料模板.xlsx').write_bytes(content)


def test_completed_contracts_do_not_hide_outstanding_or_original_expected_allocations():
    from types import SimpleNamespace as R
    from app.services.indonesia_warehouse_import import candidates
    mat = R(id='m', data={'unit': 'PCS'})
    orders, allocations = [], []
    for i in range(202):
        orders.append(R(id=str(i), version=1, data=dict(customer='C', recipient='RRI', contract_no=str(i), item_no='I',
            delivery_date='2026-12-01', plan={'required_date': '2026-11-01'}, demands=[dict(id='d', material_id='m', unit='PCS', quantity='100')])))
        if i != 201:
            allocations.append(R(data=dict(order_id=str(i), demand_id='d', normal_quantity='100', spare_quantity='0', supply_id='expected' if i == 200 else 'received')))
    supplies = [R(id='expected', data=dict(material_id='m', status='expected', purchase_no='P'))]
    ledger = R(rows=lambda kind: {'order': orders, 'allocation': allocations, 'supply': supplies}[kind])
    imported = R(data=dict(customer='C', recipient='RRI'))
    targets = candidates(ledger, imported, {'purchase_no': 'P'}, mat)
    assert [(t['order_id'], t['remaining']) for t in targets] == [('200', '0'), ('201', '100')]
    assert [t['order_id'] for t in candidates(ledger, imported, {'purchase_no': 'DIFFERENT'}, mat)] == ['201']
