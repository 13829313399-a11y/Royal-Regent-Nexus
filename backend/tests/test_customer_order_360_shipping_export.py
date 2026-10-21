"""Shipping identifiers remain separate from the PO identity in real export paths."""
from io import BytesIO

import openpyxl
import pytest

from app.schemas.customer_order import CustomerOrderLineOut
from app.services import customer_order_unified as service
from app.services.customer_order_360_workbook import organize_workbook
from test_customer_order_360_version import kwargs, parsed_row, workbook_bytes
from test_customer_order_huakang_a_unified import template
from customer_order_360_private_samples import private_po_samples


def stub(monkeypatch, product='76543220069'):
    def parse(*_):
        row = parsed_row()
        row.update(product_no=product, order_type='正单')
        return [row], [], 'TEST'
    monkeypatch.setattr(service, '_parse_huakang_a_rows', parse)


def item_row(sheet, reference='RL-765431-23'):
    return next(r for r in range(4, sheet.max_row + 1) if sheet.cell(r, 5).value == reference)


@pytest.mark.parametrize('occupied', [False, True])
def test_export_keeps_product_identity_and_reuses_shipping_column(monkeypatch, occupied):
    stub(monkeypatch)
    wb = openpyxl.load_workbook(BytesIO(template('360')))
    sheet = wb['ITEM表']
    if occupied:
        sheet['AM3'], sheet['AM4'] = '人工扩展', '版本及人工内容保留'
    source = workbook_bytes(wb)
    preview = service.create_unified_customer_preview(**kwargs(source))
    assert CustomerOrderLineOut.model_validate(preview['rows'][0]).shipping_item == '7654322'
    output, _, resolved = service.export_unified_customer_schedule(**kwargs(source))
    result = openpyxl.load_workbook(BytesIO(output))
    item = result['ITEM表']
    column = 40 if occupied else 39
    assert item.cell(3, column).value == '出货货号'
    n = item_row(item)
    assert item.cell(n, 8).value == resolved['rows'][0]['product_no'] == '76543220069'
    assert item.cell(n, column).value == '7654322'
    assert item.cell(n, column).data_type == 's'
    assert item.cell(n, column).number_format == '@'
    if occupied:
        assert item['AM3'].value == '人工扩展'
        assert item['AM4'].value == '版本及人工内容保留'
    history = service.read_unified_history(output, factory_id='huakang-a', customer_code='360')
    assert next(h for h in history if h.reference_no == 'RL-765431-23').extras['shipping_item'] == '7654322'
    result.close()


def test_manual_shipping_item_is_literal_and_does_not_change_raw_item(monkeypatch):
    stub(monkeypatch, product='00123450067')
    source = template('360')
    preview = service.create_unified_customer_preview(**kwargs(source))
    row = preview['rows'][0]
    issue = next(i for i in row['issues'] if i['field'] == 'shipping_item')
    value = '=0012345 / 另一出货号'
    output, _, resolved = service.export_unified_customer_schedule(**kwargs(source),
        manual_overrides=[dict(row_id=row['id'], issue_key=issue['skip_key'], field='shipping_item', value=value)])
    assert resolved['rows'][0]['shipping_item'] == value
    result = openpyxl.load_workbook(BytesIO(output))
    sheet = result['ITEM表']; n = item_row(sheet)
    assert sheet.cell(n, 8).value == '00123450067'
    assert sheet.cell(n, 39).value == value and sheet.cell(n, 39).data_type == 's'
    result.close()


def test_standalone_organizer_preserves_version_and_manual_shipping_across_roundtrips():
    wb = openpyxl.load_workbook(BytesIO(template('360')))
    item = wb['ITEM表']
    for cell in item[3]:
        if cell.value == '版本':
            cell.value = '其他备注'
    item['AM3'], item['AM4'] = '版本', '两个装\n业务备注'
    item['AN3'], item['AN4'] = '出货货号', '手工特别货号'
    item['E4'], item['H4'], item['Y4'] = 'RL-123456-1', '76543220069', '2026-01-01'
    for name in ('接单表', '正单评审表'):
        wb[name]['D4'], wb[name]['G4'], wb[name]['R4'] = 'RL-123456-1', '76543220069', '2026-01-01'
    first, _, _, _ = organize_workbook('schedule.xlsx', workbook_bytes(wb))
    second, _, _, _ = organize_workbook('schedule.xlsx', first)
    result = openpyxl.load_workbook(BytesIO(second))
    item = result['ITEM表']; n = item_row(item, 'RL-123456-1')
    assert item.cell(n, 39).value == '两个装\n业务备注'
    assert item.cell(n, 40).value == '手工特别货号'
    assert item.cell(n, 8).value == '76543220069'
    assert [c.column for c in item[3] if c.value == '出货货号'] == [40]
    result.close()


@pytest.mark.parametrize('sample', private_po_samples())
def test_supplied_real_po_keeps_complete_item_and_shipping_display(sample):
    po, expected, manifest = sample['source'], sample['expected'], sample['manifest']
    schedule = manifest['schedule']
    po_before, schedule_before = po.read_bytes(), schedule.read_bytes()
    preview = service.create_unified_customer_preview(customer_code='360', factory_id='huakang-a',
        received_date=manifest['received_date'], po_files=[(po.name, po_before)],
        schedule_file_name=schedule.name, schedule_content=schedule_before)
    assert len(preview['rows']) == 1
    row = CustomerOrderLineOut.model_validate(preview['rows'][0])
    assert row.reference_no == expected['reference_no']
    assert row.product_no == expected['product_no']
    assert row.shipping_item == expected['shipping_item']
    assert po.read_bytes() == po_before and schedule.read_bytes() == schedule_before
