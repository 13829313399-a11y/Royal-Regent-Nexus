from copy import deepcopy
from decimal import Decimal
from types import SimpleNamespace
import json

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from app.services.internal_quote_artifacts import _merge_electronic_quote, _clear_import_generated_payload
from app.services.internal_quote_calculator import calculate_section, CalculationInputError
from app.services.internal_quote_excel import _build_electronic_sheet, _purchase_source_lines

SNAPSHOT = {"fx": {"rmb_hkd": "0.85", "hkd_usd": "7.8"}}


def quoted(name, price, profit, labor=0):
    return {"id": name, "name": name, "pricing_currency": "RMB", "profit_rate_percent": profit,
            "labor_rmb": labor, "components": [{"item": "同名零件", "quantity": 2, "unit_price_rmb": price, "tax_rate_percent": 13}]}


def calculate(payload):
    return calculate_section("electronic", payload, SNAPSHOT, "test-ref")


def test_two_quotes_keep_independent_profit_and_sum_once():
    first, second = quoted("主板", 10, 5, 2), quoted("副板", 4, 30, 7)
    result = calculate({"quote_groups": [first, second]})
    expected = [calculate(first), calculate(second)]
    for currency in ("hkd", "rmb"):
        assert Decimal(result["totals"][f"total_{currency}"]) == sum(Decimal(item["totals"][f"total_{currency}"]) for item in expected)
    assert [g["totals"] for g in result["quote_groups"]] == [g["totals"] for g in expected]
    assert [Decimal(line["quote_pricing_hkd"]) for line in result["line_breakdown"]] == [Decimal(g["totals"]["total_hkd"]) for g in expected]


def test_mixed_legacy_hkd_and_quick_quotes_preserve_parts_for_export():
    first = {"id": "old", "name": "原报价", "components": [{"item": "旧件", "quantity": 3, "unit_price_hkd": 2}]}
    second = {"id": "quick", "name": "快捷报价", "quote_mode": "quick", "quick_quotes": [{"item": "模块", "unit_price_rmb": 10}]}
    payload = {"quote_groups": [first, second]}
    result = calculate(payload)
    section = SimpleNamespace(department="electronic", payload_json=json.dumps(payload), calculation_json=json.dumps(result))
    lines = _purchase_source_lines(section)
    assert [line["quantity"] for line in lines] == [3, 1]
    assert [line["source_unit_price"] for line in lines] == [2, 10]


def test_import_adds_then_replaces_only_selected_quote_and_rejects_repeat():
    old = quoted("unused", 10, 5)
    old.pop("id")
    old.pop("name")
    incoming = quoted("ignored", 2, 20)
    original = deepcopy(old)
    batch = SimpleNamespace(id="batch-one", source_sha256="hash-one", source_file_name="第二份.xlsx")
    added = _merge_electronic_quote(old, incoming, "new", batch, ".85")
    assert added["quote_groups"][0] == {**original, "id": "legacy", "name": "电子报价1"}
    with pytest.raises(HTTPException) as error:
        _merge_electronic_quote(added, incoming, "new", batch, ".85")
    assert error.value.status_code == 409
    target = added["quote_groups"][1]["id"]
    replaced = _merge_electronic_quote(added, quoted("ignore", 3, 25), target,
        SimpleNamespace(id="batch-two", source_sha256="hash-two", source_file_name="更新.xlsx"), ".85")
    assert replaced["quote_groups"][0] == added["quote_groups"][0]
    assert replaced["quote_groups"][1]["components"][0]["unit_price_rmb"] == 3
    assert old == original
    with pytest.raises(HTTPException):
        _merge_electronic_quote(added, incoming, None, batch, ".85")
    with pytest.raises(HTTPException):
        _merge_electronic_quote(added, incoming, "missing", batch, ".85")


def test_export_has_two_independent_electronic_sheets():
    groups = [quoted("主板电子报价", 10, 5), quoted("副板电子报价", 4, 30)]
    payload = {"quote_groups": groups}
    calculation = calculate(payload)
    section = SimpleNamespace(payload_json=json.dumps(payload), calculation_json=json.dumps(calculation))
    book = Workbook()
    _build_electronic_sheet(book, section, SNAPSHOT)
    assert book.sheetnames == ["Sheet", "电子明细1", "电子明细2"]
    for index, group in enumerate(groups, 1):
        sheet = book[f"电子明细{index}"]
        assert sheet["A1"].value == group["name"]
        assert sheet["E4"].value == group["components"][0]["unit_price_rmb"]
        total = next(row for row in sheet.iter_rows(values_only=True) if row[0] == "含税报价")
        assert total[2] == float(calculation["quote_groups"][index - 1]["totals"]["total_hkd"])


def test_deleting_source_clears_only_its_own_costs_even_when_values_match():
    groups = [quoted("first", 1, 10, 5), quoted("second", 2, 10, 5)]
    batches = []
    for group in groups:
        group["import_batch_id"] = group["id"]
        group["components"][0]["import_batch_id"] = group["id"]
        batches.append(SimpleNamespace(id=group["id"], import_type="electronic", confirmed_revision=1,
            preview_json=json.dumps({"electronic_quote_id": group["id"], "payload_fragment": {"labor_rmb": 5}})))
    cleared, _ = _clear_import_generated_payload({"quote_groups": groups}, batches[:1], all_batches=batches)
    assert cleared["quote_groups"][0]["components"] == []
    assert "labor_rmb" not in cleared["quote_groups"][0]
    assert cleared["quote_groups"][1] == groups[1]


@pytest.mark.parametrize("groups", [[quoted("same", 1, 10)] * 2, [{"id": "x", "name": "x", "quote_groups": []}], "bad"])
def test_invalid_groups_fail_visibly(groups):
    with pytest.raises(CalculationInputError):
        calculate({"quote_groups": groups})


def test_product_summary_prices_each_quote_with_its_own_margin():
    from app.services.internal_quote import _rr2_cost_summary
    from app.services.internal_quote_excel import _purchase_pricing_entries, _pricing_entry_amount
    from test_internal_quote_purchase_export import arithmetic_value
    groups = [quoted("主板", 10, 5, 2), quoted("副板", 4, 30, 7)]
    for group in groups:
        group["components"][0]["pricing_component_id"] = group["id"]
    payload = {"quote_groups": groups}
    result = calculate(payload)
    electronic = SimpleNamespace(department="electronic", is_required=True, calculation_status="valid",
        payload_json=json.dumps(payload), calculation_json=json.dumps(result))
    total = Decimal(result["totals"]["total_hkd"])
    summary = _rr2_cost_summary([electronic], {"electronic_hkd": total, "factory_price_hkd": total}, SNAPSHOT)
    entries = summary["shipping_pricing"]["pricing_entries"]
    for group, calculation in zip(groups, result["quote_groups"]):
        amount = sum(Decimal(str(entry["amount_hkd"])) for entry in entries if entry.get("electronic_quote_id") == group["id"])
        assert abs(amount - Decimal(calculation["totals"]["total_hkd"])) < Decimal("0.0001")
    book = Workbook()
    sheet = book.active
    sheet["L4"] = .85
    for index, entry in enumerate(_purchase_pricing_entries(entries, [electronic]), 1):
        sheet[f"A{index}"] = _pricing_entry_amount(entry)
        assert arithmetic_value(sheet, f"A{index}") == pytest.approx(float(entry["amount_hkd"]), abs=.0001)


def test_partial_history_selects_and_allocates_each_quote_independently():
    from app.services.internal_quote_history import _select_electronic_groups
    source = SimpleNamespace(components=[{"id": "main"}, {"id": "accessory"}])
    first, second = quoted("主板", 10, 5, 4), quoted("副板", 3, 30, 8)
    first["components"][0]["pricing_component_id"] = "main"
    first["components"].append({"item": "附件", "quantity": 2, "unit_price_rmb": 10, "pricing_component_id": "accessory"})
    second["components"][0]["pricing_component_id"] = "accessory"
    picked = _select_electronic_groups({"quote_groups": [first, second]}, source, "accessory", "new-part", Decimal(".85"), "0")
    assert [Decimal(group["labor_rmb"]) for group in picked] == [Decimal(2), Decimal(8)]
    assert [group["profit_rate_percent"] for group in picked] == [5, 30]
    assert all(row["pricing_component_id"] == "new-part" for group in picked for row in group["components"])
