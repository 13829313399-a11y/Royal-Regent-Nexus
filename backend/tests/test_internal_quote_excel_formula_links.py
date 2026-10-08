import json
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.services.internal_quote_excel import _build_summary_sheet
from app.services.internal_quote_calculator import calculate_section


def _section(
    department: str,
    *,
    payload: dict[str, object] | None = None,
    calculation: dict[str, object] | None = None,
):
    return SimpleNamespace(
        department=department,
        calculation_status="valid",
        payload_json=json.dumps(payload or {}, ensure_ascii=False),
        calculation_json=json.dumps(calculation or {}, ensure_ascii=False),
    )


def _find_row(sheet, column: int, value: object) -> int:
    return next(
        row
        for row in range(1, sheet.max_row + 1)
        if sheet.cell(row, column).value == value
    )


@pytest.mark.parametrize("mode", ["ordinary", "legacy", "color_box", "product", "pdq"])
@pytest.mark.parametrize("color_box_unit", ["inch", "cm"])
@pytest.mark.parametrize("unit,carton_count", [("inch", 1), ("inch", 2), ("cm", 2), ("inch", 3)])
def test_carton_export_uses_each_packing_quantity_and_counts_flat_cards_once(mode, color_box_unit, unit, carton_count):
    component_mode = mode != "ordinary"
    snapshot = {"fx": {"rmb_hkd": ".85", "hkd_usd": "7.8"}, "paper_price_factor": "2.75"}
    cartons = [
        {"item": "主纸箱", "size_unit": unit, "length_in": 18, "width_in": 12, "height_in": 10, "qty_per_carton": 24,
         "flat_cards": [{"name": "外箱平卡", "length_in": 10, "width_in": 4, "quantity": 2}]},
        {"item": "内纸箱", "size_unit": unit, "length_in": 5, "width_in": 5, "height_in": 4, "qty_per_carton": 6,
         "flat_cards": [{"name": "内箱平卡", "length_in": 3, "width_in": 2, "quantity": 1}]},
        {"item": "内纸箱2", "size_unit": unit, "length_in": 8, "width_in": 6, "height_in": 5, "qty_per_carton": 12},
    ][:carton_count]
    payload = {"cartons": cartons, "paper_price_factor": 2.75, "inner_paper_price_factor": 3.25,
               "color_box_size_in": {"length": 17.25, "width": 11.25, "height": 9}, "color_box_size_unit": color_box_unit,
               "testing_fee_enabled": False, "freight_calc": {"enabled": False}}
    if component_mode:
        payload["pricing_mode"] = "component"
    if mode in {"color_box", "product", "pdq"}:
        payload["justplay_carton"] = {"dimension_source": "product" if mode == "product" else "color_box", "length_count": 2, "width_count": 3, "height_count": 4}
        payload["color_box_size_in"] = {"length": 8.625, "width": 3.75, "height": 2.25}
        payload["product_size_in"] = {"length": 8.625, "width": 3.75, "height": 2.25}
    if mode == "pdq":
        payload["pdq_size_in"] = {"length": 17.25, "width": 11.25, "height": 9}
        payload["pdq_size_unit"] = color_box_unit
    calculation = calculate_section("sales", payload, snapshot, "TEST-CARTON")
    # Independently calculated from the screenshot dimensions and each carton quantity.
    expected = .1686666666666667 + .22 / 24
    if carton_count >= 2:
        expected += .13 + .0165 / 6
    if carton_count == 3:
        expected += .104
    assert float(calculation["totals"]["carton_hkd"]) == pytest.approx(expected, abs=.00005)
    entries = [{"section": "sales", "is_global": True, "label": row["item"], "amount_hkd": row.get("per_piece_hkd", row.get("amount_hkd")), **row}
               for row in calculation["line_breakdown"]]
    shipping = {"markup": "1.2", "misc_ratio": ".02", "rows": [], "freight_enabled": False, "lifting_enabled": False}
    if component_mode:
        shipping.update(pricing_mode="component", pricing_groups=[{"id": "main", "name": "主体", "markup": "1.15", "cost_hkd": "0"}],
                        pricing_entries=entries, global_pricing={"entries": entries, "markup": "1.25", "cost_hkd": calculation["totals"]["total_hkd"]})
    workbook = Workbook()
    try:
        _build_summary_sheet(workbook, SimpleNamespace(product_name="内外箱验证", quote_no="CARTON-TEST", customer="JustPlay" if component_mode else "普通客", region_code="mainland", remark=""),
                             [_section("sales", payload=payload, calculation=calculation)], snapshot,
                             {"shipping_pricing": shipping, "t1": [], "t2": [], "t3": [], "t4": []},
                             {"factory_price_hkd": calculation["totals"]["total_hkd"], "carton_hkd": calculation["totals"]["carton_hkd"]}, [])
        sheet = workbook["报价明细"]
        outer = _find_row(sheet, 14, f"外箱 ({unit}):")
        packing = _find_row(sheet, 14, "装箱：")
        total = _find_row(sheet, 14, "合计")
        board = _find_row(sheet, 14, "纸板价")
        color_box = _find_row(sheet, 14, f"彩盒尺寸 ({color_box_unit})")
        for column, base, offset in ((15, 17.25, ".75"), (16, 11.25, ".75"), (17, 9, "1")):
            letter = chr(ord('A') + column - 1)
            count = [2, 3, 4][column - 15] if mode in {"color_box", "product", "pdq"} else 1
            assert sheet.cell(color_box, column).value == pytest.approx(base / count * (2.54 if color_box_unit == "cm" else 1))
            if component_mode:
                source_row = (_find_row(sheet, 14, f"PDQ 尺寸 ({color_box_unit})") if mode == "pdq" else
                              _find_row(sheet, 14, "产品尺寸 (in)") if mode == "product" else color_box)
                source_ref = f"{letter}{source_row}" + ("/2.54" if color_box_unit == "cm" and mode != "product" else "")
                count_ref = "" if mode == "pdq" else f"*{letter}{_find_row(sheet, 14, '方向个数')}"
                expected_formula = f"=IF({letter}{source_row}>0,({source_ref}{count_ref}+{float(offset):g})" + ("*2.54" if unit == "cm" else "") + ",0)"
                assert sheet.cell(outer, column).value == expected_formula
            else:
                assert sheet.cell(outer, column).value == pytest.approx((base + float(offset)) * (2.54 if unit == "cm" else 1))
        assert sheet.cell(packing, 15).value == ([24, "6/24", "6/12/24"][carton_count - 1])
        if mode == "pdq":
            assert _find_row(sheet, 14, f"PDQ 尺寸 ({color_box_unit})") == color_box + 1
            assert not any(cell.value == "方向个数" for cell in sheet['N'])
        else:
            assert not any(str(cell.value).startswith("PDQ 尺寸") for cell in sheet['N'])
        assert sheet.cell(total, 15).value.startswith("=SUM(" if carton_count > 1 else "=IFERROR(")
        assert f"O{board}" in sheet.cell(total, 15).value
        price_labels = ["箱价："] if carton_count == 1 else ["外箱价：", "内箱价：", "内箱2价："][:carton_count]
        for index, label in enumerate(price_labels):
            price = _find_row(sheet, 14, label)
            dimension_row = outer + index
            formula = sheet.cell(price, 15).value
            assert f"O{dimension_row}" in formula and f"P{dimension_row}" in formula and f"Q{dimension_row}" in formula
            assert ("$N$6" if index == 0 else "$O$6") in formula
            assert "_报价明细计算" not in formula
            if carton_count > 1:
                assert f"O{packing}" in formula
                detail_label = ["外纸箱", "内纸箱", "内纸箱2"][index]
                assert sheet.cell(_find_row(sheet, 3, detail_label), 4).value == f"=O{price}"
        if carton_count > 1:
            assert _find_row(sheet, 3, "内纸箱") < _find_row(sheet, 3, "外纸箱")
            assert sheet.cell(_find_row(sheet, 3, "平卡"), 4).value == f"=O{board}"
        if component_mode:
            adhesive = sheet.cell(_find_row(sheet, 3, "胶纸/胶水/胶针"), 4).value
            assert f"O{packing}" in adhesive
            if unit == "cm":
                assert f"O{outer}/2.54" in adhesive
            pallet_formula = sheet.cell(_find_row(sheet, 3, "纸托板成本"), 4).value
            assert pallet_formula.startswith("=19/30/")
            assert f"O{packing}" in pallet_formula
            assert pallet_formula.endswith("+0.0")
            # Both costs must divide by the outer carton quantity, including "inner/outer" layouts.
            if carton_count == 1:
                quantity_ref = f"O{packing}"
            elif carton_count == 2:
                quantity_ref = f'VALUE(RIGHT(O{packing},LEN(O{packing})-FIND("/",O{packing})))'
            else:
                quantity_ref = f'VALUE(TRIM(MID(SUBSTITUTE(O{packing},"/",REPT(" ",32)),65,32)))'
            assert adhesive.endswith(f"/{quantity_ref}+0.0")
            assert pallet_formula == f"=19/30/{quantity_ref}+0.0"
            pallet = next(row for row in calculation["line_breakdown"] if row.get("formula_code") == "paper_pallet")
            assert pallet["cartons_per_pallet"] == "30.0000"
        for row in range(outer, total + 1):
            assert sheet.cell(row, 14).border.left.style
            assert sheet.cell(row, 17).border.right.style
    finally:
        workbook.close()


@pytest.mark.parametrize("detached_index", [0, 1])
def test_carton_override_keeps_live_carton_and_own_flat_card_formulas(detached_index):
    cartons = [
        {"item": "主纸箱", "length_in": 18, "width_in": 12, "height_in": 10, "qty_per_carton": 24,
         "flat_cards": [{"length_in": 10, "width_in": 4, "quantity": 2}]},
        {"item": "内纸箱", "length_in": 5, "width_in": 5, "height_in": 4, "qty_per_carton": 6,
         "flat_cards": [{"length_in": 3, "width_in": 2, "quantity": 1}]},
    ]
    cartons[detached_index]["markup_override"] = "1.08"
    payload = {"cartons": cartons, "paper_price_factor": 2.75, "inner_paper_price_factor": 3.25,
               "testing_fee_enabled": False, "freight_calc": {"enabled": False}}
    snapshot = {"fx": {"rmb_hkd": ".85", "hkd_usd": "7.8"}}
    calculation = calculate_section("sales", payload, snapshot, "TEST-CARTON-OVERRIDE")
    line = calculation["line_breakdown"][detached_index]
    entry = {**line, "section": "sales", "label": line["item"], "amount_hkd": line["per_piece_hkd"]}
    summary = {"t1": [], "t2": [], "t3": [], "t4": [], "shipping_pricing": {
        "markup": "1.2", "misc_ratio": ".02", "rows": [], "pricing_mode": "standard",
        "pricing_entries": [entry], "pricing_groups": [
            {"id": "main", "name": "主倍率汇总", "markup": "1.2", "cost_hkd": "0"},
            {"id": "detail-01", "name": line["item"], "markup": "1.08", "cost_hkd": line["per_piece_hkd"]},
        ],
    }}
    workbook = Workbook()
    try:
        _build_summary_sheet(workbook, SimpleNamespace(product_name="内外箱独立倍率", quote_no="CARTON-SPLIT", region_code="mainland", remark=""),
                             [_section("sales", payload=payload, calculation=calculation)], snapshot, summary,
                             {"carton_hkd": calculation["totals"]["carton_hkd"]})
        sheet = workbook["报价明细"]
        label = "外纸箱" if detached_index == 0 else "内纸箱"
        detail_rows = [row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 3).value == label]
        assert len(detail_rows) == 2
        assert sheet.cell(detail_rows[0], 4).value == "=0"
        price = _find_row(sheet, 14, "外箱价：" if detached_index == 0 else "内箱价：")
        detached_formula = sheet.cell(detail_rows[1], 4).value
        assert detached_formula.startswith(f"=O{price}+(")
        packing = _find_row(sheet, 14, "装箱：")
        assert f"O{packing}" in detached_formula
        own_board_formula = detached_formula.split("+(", 1)[1][:-1]
        main_board = sheet.cell(_find_row(sheet, 3, "平卡"), 4).value
        assert main_board.endswith(f"-({own_board_formula})")
        assert sheet.cell(detail_rows[1] + 2, 4).value == 1.08
    finally:
        workbook.close()


def test_export_links_molding_assembly_cartons_and_flat_cards_to_source_cells():
    workbook = Workbook()
    quote = SimpleNamespace(
        product_name="公式联动测试",
        quote_no="IQ-FORMULA-LINKS",
        region_code="mainland",
        remark="",
    )
    sales = _section(
        "sales",
        payload={
            "paper_price_factor": 2.75,
            "inner_paper_price_factor": 3.25,
            "flat_card_price_factor": 2.75,
            "color_box_size_unit": "cm",
            "color_box_size_in": {"length": 7.5, "width": 3.35, "height": 4.92},
            "product_size_in": {"length": 6, "width": 6, "height": 6},
            "cartons": [
                {
                    "item": "主纸箱",
                    "size_unit": "cm",
                    "length_in": 80 / 2.54,
                    "width_in": 29 / 2.54,
                    "height_in": 30 / 2.54,
                    "qty_per_carton": 4,
                    "flat_cards": [
                        {
                            "name": "主平卡",
                            "length_in": 19 / 2.54,
                            "width_in": 8.5 / 2.54,
                            "quantity": 2,
                        }
                    ],
                },
                {
                    "item": "内纸箱",
                    "size_unit": "cm",
                    "length_in": 80 / 2.54,
                    "width_in": 29 / 2.54,
                    "height_in": 30 / 2.54,
                    "qty_per_carton": 1,
                    "flat_cards": [],
                },
            ],
        },
        calculation={"line_breakdown": [], "totals": {}},
    )
    molding = _section(
        "molding",
        calculation={
            "line_breakdown": [
                {
                    "kind": "injection",
                    "item": "车轮芯",
                    "material": "1#PP",
                    "grade": "7032 E3",
                    "net_weight_g": "3.2",
                    "loss_rate_percent": "3",
                    "loss_weight_g": "3.296",
                    "material_price_hkd_lb": "6.8",
                    "material_price_hkd_g": "0.014978",
                    "machine_code": "4A-6A",
                    "machine_shift_price_hkd": "940",
                    "sets": "3",
                    "target_output": "4800",
                    "molding_cost_hkd": "0.065278",
                    "material_cost_hkd": "0.04937",
                }
            ],
            "totals": {},
        },
    )
    assembly_amount = 35 * 260 / 3500
    packaging_amount = 17 * 260 / 3500
    assembly = _section(
        "assembly",
        payload={"labor_base_hkd": 260},
        calculation={
            "group_summaries": [
                {
                    "category": "assembly",
                    "group": "组装半成品",
                    "standard_work_hours": "11",
                    "labor_base_hkd": "260",
                    "production_qty": "3500",
                    "teams": "1",
                    "total_persons": "35",
                    "amount_hkd_pcs": str(assembly_amount),
                },
                {
                    "category": "packaging",
                    "group": "包装",
                    "standard_work_hours": "11",
                    "labor_base_hkd": "260",
                    "production_qty": "3500",
                    "teams": "1",
                    "total_persons": "17",
                    "amount_hkd_pcs": str(packaging_amount),
                },
            ]
        },
    )
    summary = {
        "t1": [{"key": "material", "value": "0.04937"}],
        "t2": [],
        "t3": [{"key": "injection_labor", "value": "0.065278"}],
        "t4": [],
        "molding_material_breakdown": {
            "total_hkd": "0.04937",
            "imported_hkd": "0.04937",
            "domestic_hkd": "0",
        },
        "shipping_pricing": {
            "markup": "1.2",
            "active_markup_moq": "10000",
            "markup_tiers": [],
            "misc_ratio": "0.02",
            "rows": [],
        },
    }
    reference_snapshot = {
        "fx": {"rmb_hkd": "0.85", "hkd_usd": "7.8"},
        "material_prices": {
            "1#PP|7032 E3": {
                "label": "7032 E3 1#PP料",
                "price_hkd_lb": "6.8",
            }
        },
        "machine_prices": [
            {"machine_range": "4A-6A", "shift_price_hkd": "940"}
        ],
        "assembly_labor_base_hkd": "260",
    }
    _build_summary_sheet(
        workbook,
        quote,
        [sales, molding, assembly],
        reference_snapshot,
        summary,
        {
            "molding_hkd": 0.04937 + 0.065278,
            "assembly_hkd": assembly_amount,
            "packing_labor_hkd": packaging_amount,
        },
    )
    sheet = workbook.active
    try:
        assert sheet["E9"].value == "=3.2*(1+3/100)"
        assert sheet["F9"].value == "=D$2/454"
        assert sheet["J9"].value == "=D$6/I9/H9"
        assert sheet["K9"].value == "=E9*F9"

        molding_total_row = _find_row(sheet, 3, "模具合计")
        material_detail_row = _find_row(sheet, 2, "料价")
        molding_detail_row = _find_row(sheet, 2, "啤工")
        assert sheet.cell(material_detail_row, 4).value == f"=K{molding_total_row}"
        assert sheet.cell(molding_detail_row, 4).value == f"=J{molding_total_row}"

        assembly_row = next(
            row
            for row in range(1, sheet.max_row + 1)
            if str(sheet.cell(row, 3).value).startswith("组装半成品")
        )
        packaging_labor_row = next(
            row
            for row in range(1, sheet.max_row + 1)
            if str(sheet.cell(row, 3).value).startswith("包装（")
        )
        assert sheet.cell(assembly_row, 4).value == "=35*$M$4/3500"
        assert sheet.cell(packaging_labor_row, 4).value == "=17*$M$4/3500"

        assert [sheet.cell(5, column).value for column in range(14, 18)] == [
            "主纸箱系数",
            "内纸箱系数",
            "平卡系数",
            "杂项",
        ]
        assert [sheet.cell(6, column).value for column in range(14, 17)] == [
            2.75,
            3.25,
            2.75,
        ]

        outer_row = _find_row(sheet, 14, "外箱 (cm):")
        inner_row = _find_row(sheet, 14, "内箱 (cm):")
        flat_card_row = _find_row(sheet, 14, "主平卡 (cm):")
        paperboard_row = _find_row(sheet, 14, "纸板价")
        carton_price_row = _find_row(sheet, 14, "外箱价：")
        inner_price_row = _find_row(sheet, 14, "内箱价：")
        packing_qty_row = _find_row(sheet, 14, "装箱：")
        total_row = _find_row(sheet, 14, "合计")
        assert inner_row == outer_row + 1
        assert sheet.cell(packing_qty_row, 15).value == "1/4"
        outer_quantity = f'VALUE(RIGHT(O{packing_qty_row},LEN(O{packing_qty_row})-FIND("/",O{packing_qty_row})))'
        inner_quantity = f'VALUE(LEFT(O{packing_qty_row},FIND("/",O{packing_qty_row})-1))'
        assert sheet.cell(flat_card_row, 17).value == 2.0
        assert sheet.cell(paperboard_row, 15).value == (
            f"=O{flat_card_row}/2.54*P{flat_card_row}/2.54*Q{flat_card_row}*$P$6/1000/{outer_quantity}"
        )
        assert sheet.cell(carton_price_row, 15).value == (
            f"=(O{outer_row}/2.54+P{outer_row}/2.54+2)"
            f"*(P{outer_row}/2.54+Q{outer_row}/2.54+1)*$N$6*2/1000/{outer_quantity}"
        )
        assert sheet.cell(inner_price_row, 15).value == (
            f"=(O{inner_row}/2.54+P{inner_row}/2.54+2)"
            f"*(P{inner_row}/2.54+Q{inner_row}/2.54+1)*$O$6*2/1000/{inner_quantity}"
        )
        assert sheet.cell(total_row, 15).value == (
            f"=SUM(O{paperboard_row}:O{inner_price_row})"
        )
        assert sheet.cell(_find_row(sheet, 3, "内纸箱"), 4).value == f"=O{inner_price_row}"
        assert sheet.cell(_find_row(sheet, 3, "外纸箱"), 4).value == f"=O{carton_price_row}"
        assert sheet.cell(_find_row(sheet, 3, "平卡"), 4).value == f"=O{paperboard_row}"
    finally:
        workbook.close()


def test_export_renders_formula_driven_split_pricing_groups():
    workbook = Workbook()
    quote = SimpleNamespace(product_name="分项倍率测试", quote_no="IQ-SPLIT-PRICING", region_code="mainland", remark="")
    sales = _section("sales", payload={"testing_fee_enabled": False}, calculation={"line_breakdown": [], "totals": {}})
    summary = {
        "t1": [
            {"key": "hardware", "value": "2.13"},
            {"key": "sewing_cloth", "value": "12.66"},
        ],
        "t2": [{"key": "other_buy", "value": "4.2"}],
        "t3": [],
        "t4": [],
        "shipping_pricing": {
            "markup": "1.16",
            "active_markup_moq": "10000",
            "markup_tiers": [],
            "misc_ratio": "0.02",
            "pricing_mode": "standard",
            "pricing_groups": [
                {"id": "main", "name": "主倍率汇总", "cost_hkd": "2.13", "pricing_base_hkd": "2.13", "markup": "1.16"},
                {"id": "detail-01", "name": "车衣", "cost_hkd": "12.66", "pricing_base_hkd": "12.66", "markup": "1.03"},
                {"id": "detail-02", "name": "镜片", "cost_hkd": "4.2", "pricing_base_hkd": "4.2", "markup": "1.08"},
            ],
            "pricing_entries": [
                {
                    "section": "sewing",
                    "kind": "sewing",
                    "label": "车衣",
                    "amount_hkd": "12.66",
                    "markup_override": "1.03",
                },
                {
                    "section": "engineering",
                    "kind": "purchase",
                    "auxiliary_category": "镜片",
                    "label": "镜片",
                    "amount_hkd": "4.2",
                    "markup_override": "1.08",
                },
            ],
            "rows": [
                {"name": "出厂价", "freight_hkd": "0", "lift_hkd": "0"},
                {"name": "港柜", "freight_hkd": "1.90", "lift_hkd": "1.10"},
            ],
        },
    }
    _build_summary_sheet(
        workbook,
        quote,
        [sales],
        {"fx": {"rmb_hkd": "0.85", "hkd_usd": "7.8"}},
        summary,
        {"sewing_hkd": "12.66"},
    )
    sheet = workbook.active
    try:
        def group_cost_row(name: str) -> int:
            return next(
                row
                for row in range(1, sheet.max_row + 1)
                if sheet.cell(row, 2).value == name
                and sheet.cell(row, 3).value == "成本"
            )

        main_cost_row = group_cost_row("主倍率汇总")
        detail_cost_row = group_cost_row("车衣")
        mirror_cost_row = group_cost_row("镜片")
        total_quote_row = _find_row(sheet, 2, "报价合计")
        main_cost_formula = str(sheet.cell(main_cost_row, 4).value)
        assert main_cost_formula.startswith("=SUM(D")
        assert main_cost_formula != f"=D{main_cost_row}"
        assert str(sheet.cell(main_cost_row, 5).value).startswith(
            f"=$D${main_cost_row}+"
        )
        assert sheet.cell(main_cost_row + 1, 4).value == 1.16
        assert sheet.cell(main_cost_row + 2, 4).value == "=1-$Q$6"
        assert sheet.cell(main_cost_row + 3, 4).value == f"=D{main_cost_row}*D{main_cost_row + 1}/D{main_cost_row + 2}"
        assert sheet.cell(detail_cost_row - 1, 2).value == "车衣"
        assert sheet.cell(detail_cost_row - 1, 3).value == "车衣"
        assert str(sheet.cell(detail_cost_row - 1, 4).value).startswith("=")
        assert sheet.cell(detail_cost_row + 1, 4).value == 1.03
        assert sheet.cell(mirror_cost_row - 1, 2).value == "其他外购"
        assert sheet.cell(mirror_cost_row - 1, 3).value == "镜片"
        assert sheet.cell(mirror_cost_row + 1, 4).value == 1.08
        assert detail_cost_row == main_cost_row + 6
        assert mirror_cost_row == detail_cost_row + 6
        assert total_quote_row == mirror_cost_row + 4
        assert sheet.cell(main_cost_row + 4, 2).value is None
        assert sheet.cell(detail_cost_row + 4, 2).value is None
        assert sheet.cell(total_quote_row, 4).value == (
            f"=D{main_cost_row + 3}+D{detail_cost_row + 3}+D{mirror_cost_row + 3}"
        )
        def border_style(row: int, column: int, side: str) -> str | None:
            return getattr(getattr(sheet.cell(row, column).border, side), "style", None)

        assert border_style(main_cost_row, 2, "top") is None
        assert border_style(main_cost_row, 4, "top") == "thin"
        assert border_style(main_cost_row, 4, "left") is None
        assert border_style(main_cost_row + 3, 2, "bottom") is None
        assert border_style(main_cost_row + 3, 4, "bottom") == "thin"
        assert border_style(detail_cost_row - 1, 2, "top") is None
        assert border_style(detail_cost_row - 1, 4, "top") == "thin"
        assert border_style(detail_cost_row + 3, 4, "bottom") == "thin"
        assert border_style(total_quote_row, 2, "bottom") is None
        assert border_style(total_quote_row, 4, "bottom") == "thin"
        assert not any(
            sheet.cell(row, 3).value in {
                "分段汇总（HKD）",
                "分段报价合计（HKD）",
                "综合报价（USD，最后换算）",
            }
            for row in range(1, sheet.max_row + 1)
        )
        legacy_header_row = _find_row(sheet, 3, "旺季价")
        assert sheet.cell(legacy_header_row + 1, 4).value == f"=D{total_quote_row}"
        assert str(sheet.cell(legacy_header_row + 1, 9).value).count("SUMIF") >= 3
        second_header_row = legacy_header_row + 3
        assert str(sheet.cell(second_header_row + 1, 9).value).count("SUMIF") >= 3
        third_header_row = legacy_header_row + 6
        assert sheet.cell(third_header_row, 15).value in (None, "")
        assert sheet.cell(third_header_row + 1, 15).value in (None, "")
        tax_amount_row = legacy_header_row + 10
        assert str(sheet.cell(tax_amount_row, 4).value).count("SUMIF") >= 3
        assert not any(
            "#REF!" in str(cell.value)
            for row in sheet.iter_rows()
            for cell in row
        )
    finally:
        workbook.close()


@pytest.mark.parametrize("parameters,expected_extras", [
    (None, ("0.0", "0.0")),
    ({"adhesive_extra_hkd": 0.12, "cartons_per_pallet": 40, "paper_pallet_extra_hkd": 0.07}, ("0.1", "0.1")),
    ({"adhesive_extra_hkd": 0, "cartons_per_pallet": 30, "paper_pallet_extra_hkd": 0}, ("0.0", "0.0")),
    ({"adhesive_extra_hkd": 0.062, "paper_pallet_extra_hkd": 0.049}, ("0.1", "0.0")),
    ({"adhesive_extra_hkd": 0.125, "paper_pallet_extra_hkd": 0.045}, ("0.1", "0.0")),
    ({"pallet_length_mm": 1000, "pallet_width_mm": 1200, "pallet_height_mm": 1500, "adhesive_extra_hkd": 0.15}, ("0.2", "0.0")),
    ({"cartons_per_pallet_override": 20, "paper_pallet_extra_hkd": 0.1}, ("0.0", "0.1")),
])
def test_export_renders_justplay_components_before_one_global_packaging_block(parameters, expected_extras):
    workbook = Workbook()
    quote = SimpleNamespace(
        product_name="JustPlay 分配输出测试",
        quote_no="IQ-JUSTPLAY-LAYOUT",
        customer="JustPlay",
        region_code="mainland",
        remark="",
    )
    sales = _section(
        "sales",
        payload={
            "pricing_mode": "component",
            **({"justplay_packaging": parameters} if parameters is not None else {}),
            "pricing_components": [
                {"id": "component-01", "name": "主体", "markup_x": "1.15"},
                {"id": "component-02", "name": "镜子", "markup_x": "1.05"},
            ],
            "color_box_size_in": {"length": 11.25, "width": 9.25, "height": 7},
            "testing_fee_enabled": False,
            "cartons": [
                {
                    "length_in": "12",
                    "width_in": "10",
                    "height_in": "8",
                    "size_unit": "in",
                    "qty_per_carton": "6",
                }
            ],
        },
        calculation={"line_breakdown": [], "totals": {}},
    )
    molding = _section(
        "molding",
        calculation={
            "line_breakdown": [
                {
                    "kind": "injection",
                    "item": "主体胶壳",
                    "material": "PP",
                    "net_weight_g": "97",
                    "loss_rate_percent": "3",
                    "loss_weight_g": "99.91",
                    "material_price_hkd_lb": "6.8",
                    "machine_shift_price_hkd": "940",
                    "sets": "1",
                    "target_output": "1175",
                    "material_cost_hkd": "1.2",
                    "molding_cost_hkd": "0.8",
                    "amount_hkd": "2",
                    "pricing_component_id": "component-01",
                },
                {
                    "kind": "injection",
                    "item": "镜框",
                    "material": "ABS",
                    "net_weight_g": "31",
                    "loss_rate_percent": "3",
                    "loss_weight_g": "31.93",
                    "material_price_hkd_lb": "8.5",
                    "machine_shift_price_hkd": "1490",
                    "sets": "1",
                    "target_output": "3725",
                    "material_cost_hkd": "0.6",
                    "molding_cost_hkd": "0.4",
                    "amount_hkd": "1",
                    "pricing_component_id": "component-02",
                },
            ],
            "totals": {},
        },
    )
    summary = {
        "t1": [],
        "t2": [],
        "t3": [],
        "t4": [],
        "shipping_pricing": {
            "markup": "1.15",
            "active_markup_moq": "10000",
            "markup_tiers": [],
            "misc_ratio": "0.03",
            "pricing_mode": "component",
            "pricing_groups": [
                {"id": "component-01", "name": "主体", "cost_hkd": "2.52", "pricing_base_hkd": "4.02", "markup": "1.15"},
                {"id": "component-02", "name": "镜子", "cost_hkd": "1", "pricing_base_hkd": "1", "markup": "1.05"},
            ],
            "pricing_entries": [
                {"section": "molding", "kind": "injection", "label": "主体胶壳", "amount_hkd": "2", "pricing_component_id": "component-01", "is_global": False},
                {"section": "molding", "kind": "injection", "label": "镜框", "amount_hkd": "1", "pricing_component_id": "component-02", "is_global": False},
                {
                    "section": "assembly",
                    "kind": "assembly_process",
                    "label": "主体组装",
                    "amount_hkd": "0.52",
                    "pricing_component_id": "component-01",
                    "is_global": False,
                    "persons": "7",
                    "teams": "1",
                    "production_qty": "3500",
                    "formula_allocation_factor": "1",
                },
                {
                    "section": "assembly",
                    "kind": "assembly_process",
                    "category": "packaging",
                    "label": "成品 A",
                    "amount_hkd": "0.26",
                    "is_global": True,
                    "persons": "2",
                    "teams": "1",
                    "production_qty": "2000",
                    "formula_allocation_factor": "1",
                },
                {"section": "sales", "kind": "packaging_material", "category": "color_box_inner_card", "label": "彩盒/内卡", "amount_hkd": "0.8", "is_global": True},
                {"section": "sales", "kind": "justplay_fixed_packaging", "category": "other_purchase", "label": "胶纸/胶水/胶针", "formula_code": "adhesive", "amount_hkd": "0.0811627907", "is_global": True},
                {"section": "sales", "kind": "justplay_fixed_packaging", "category": "other_purchase", "label": "纸托板成本", "formula_code": "paper_pallet", "amount_hkd": "0.1819444444", "is_global": True},
                {"section": "sales", "kind": "carton", "label": "外箱", "amount_hkd": "0.7", "is_global": True},
            ],
            "global_pricing": {
                "cost_hkd": "1.7631072351",
                "pricing_base_hkd": "1.7631072351",
                "markup": "1.25",
                "settlement": "0.97",
                "entries": [
                    {
                        "section": "assembly",
                        "kind": "assembly_process",
                        "category": "packaging",
                        "label": "成品 A",
                        "amount_hkd": "0.26",
                        "is_global": True,
                        "persons": "2",
                        "teams": "1",
                        "production_qty": "2000",
                        "formula_allocation_factor": "1",
                    },
                    {"section": "sales", "kind": "packaging_material", "category": "color_box_inner_card", "label": "彩盒/内卡", "amount_hkd": "0.8", "is_global": True},
                    {"section": "sales", "kind": "justplay_fixed_packaging", "category": "other_purchase", "label": "胶纸/胶水/胶针", "formula_code": "adhesive", "amount_hkd": "0.0811627907", "is_global": True},
                    {"section": "sales", "kind": "justplay_fixed_packaging", "category": "other_purchase", "label": "纸托板成本", "formula_code": "paper_pallet", "amount_hkd": "0.1819444444", "is_global": True},
                    {"section": "sales", "kind": "carton", "label": "外箱", "amount_hkd": "0.7", "is_global": True},
                ],
            },
            "rows": [
                {"name": "出厂价", "freight_hkd": "0", "lift_hkd": "0"},
                {"name": "港柜", "freight_hkd": "2.645", "lift_hkd": "3.657"},
                {"name": "港散", "freight_hkd": "5.447", "lift_hkd": "10.169"},
            ],
        },
    }
    _build_summary_sheet(
        workbook,
        quote,
        [sales, molding],
        {"fx": {"rmb_hkd": "0.85", "hkd_usd": "7.8"}},
        summary,
        {},
    )
    sheet = workbook["报价明细"]
    try:
        assert workbook.sheetnames[:2] == ["报价明细", "_报价明细计算"]
        main_title_row = _find_row(sheet, 1, "主体-明细")
        mirror_title_row = _find_row(sheet, 1, "镜子-明细")
        packaging_title_row = _find_row(sheet, 3, "包装明细")

        def find_between(column: int, value: str, start: int, end: int) -> int:
            return next(
                row
                for row in range(start, end + 1)
                if sheet.cell(row, column).value == value
            )

        assert main_title_row < mirror_title_row < packaging_title_row
        assert _find_row(sheet, 3, "主体胶壳") < mirror_title_row
        assert _find_row(sheet, 3, "镜框") > mirror_title_row
        main_mold_row = _find_row(sheet, 3, "主体胶壳")
        mirror_mold_row = _find_row(sheet, 3, "镜框")
        assert sheet.cell(main_mold_row, 5).value == "=97*(1+3/100)"
        assert sheet.cell(main_mold_row, 6).value.startswith("=")
        assert sheet.cell(main_mold_row, 10).value.startswith("=")
        assert sheet.cell(main_mold_row, 11).value == (
            f"=E{main_mold_row}*F{main_mold_row}"
        )
        assert sheet.cell(mirror_mold_row, 5).value == "=31*(1+3/100)"
        assembly_row = _find_row(sheet, 3, "主体组装")
        assert sheet.cell(assembly_row, 4).value == "=7*$M$4/3500"
        main_hkd_row = find_between(3, "配件报价（HKD）：", main_title_row, mirror_title_row - 1)
        main_fx_row = find_between(3, "美金兑港币汇率：", main_title_row, mirror_title_row - 1)
        main_usd_row = find_between(3, "配件单价（USD）：", main_title_row, mirror_title_row - 1)
        assert main_fx_row == main_hkd_row + 1
        assert main_usd_row == main_hkd_row + 2
        assert sheet.cell(main_fx_row, 4).value == "=$R$4"
        assert sheet.cell(main_usd_row, 4).value == f"=D{main_hkd_row}/D{main_fx_row}"
        mirror_hkd_row = find_between(3, "配件报价（HKD）：", mirror_title_row, packaging_title_row - 1)
        mirror_fx_row = find_between(3, "美金兑港币汇率：", mirror_title_row, packaging_title_row - 1)
        mirror_usd_row = find_between(3, "配件单价（USD）：", mirror_title_row, packaging_title_row - 1)
        assert mirror_fx_row == mirror_hkd_row + 1
        assert mirror_usd_row == mirror_hkd_row + 2
        assert sheet.cell(mirror_usd_row, 4).value == f"=D{mirror_hkd_row}/D{mirror_fx_row}"
        assert sum(
            1
            for row in range(1, sheet.max_row + 1)
            if sheet.cell(row, 3).value == "包装明细"
        ) == 1
        packaging_cost_row = _find_row(sheet, 3, "彩盒/内卡")
        assert packaging_cost_row > packaging_title_row
        assert sheet.cell(packaging_cost_row, 4).value == 0.8
        packaging_labor_row = _find_row(sheet, 3, "包装人工 - 成品 A")
        assert packaging_title_row < packaging_labor_row < packaging_cost_row
        assert sheet.cell(packaging_labor_row, 2).value == "装配工"
        assert sheet.cell(packaging_labor_row, 4).value == "=2*$M$4/2000"
        outer_carton_row = _find_row(sheet, 14, "外箱 (inch):")
        packing_qty_row = _find_row(sheet, 14, "装箱：")
        adhesive_row = _find_row(sheet, 3, "胶纸/胶水/胶针")
        paper_pallet_row = _find_row(sheet, 3, "纸托板成本")
        assert packaging_title_row < adhesive_row < paper_pallet_row
        assert sheet.cell(adhesive_row, 4).value == (
            f"=3.9/2150*(O{outer_carton_row}*2+P{outer_carton_row}*4+6)/O{packing_qty_row}+{expected_extras[0]}"
        )
        # 12 x 10 x 8 inch carton => 4 x 3 x 6 = 72, ignoring historical manual counts.
        expected_capacity = 84 if parameters and parameters.get("pallet_height_mm") == 1500 else 72
        if parameters and parameters.get("cartons_per_pallet_override") is not None:
            expected_capacity = parameters["cartons_per_pallet_override"]
        assert sheet.cell(paper_pallet_row, 4).value == f"=19/{expected_capacity}/O{packing_qty_row}+{expected_extras[1]}"
        for detail_row in (adhesive_row, paper_pallet_row):
            assert sheet.cell(detail_row, 4).data_type == "f"
            assert sheet.cell(detail_row, 4).number_format == "0.000"
        for exported_sheet in workbook:
            for row in exported_sheet.iter_rows():
                for cell in row:
                    assert cell.value not in (
                        "胶纸及纸托板参数", "每托板装箱数",
                        "胶纸附加金额 (HKD/件)", "纸托板附加金额 (HKD/件)",
                    )
                    assert "INT(1150/" not in str(cell.value)
        assert sheet.cell(adhesive_row, 1).value == ""
        assert sheet.cell(paper_pallet_row, 1).value == ""
        carton_cost_row = _find_row(sheet, 3, "纸箱")
        assert str(sheet.cell(carton_cost_row, 4).value).startswith("=O")
        lift_cost_row = _find_row(sheet, 3, "吊柜费")
        packaging_freight_row = find_between(
            3,
            "运费",
            packaging_title_row,
            lift_cost_row - 1,
        )
        packaging_subtotal_row = find_between(
            3,
            "成本金额：",
            packaging_freight_row + 1,
            lift_cost_row - 1,
        )
        packaging_hkd_row = find_between(
            3,
            "单价（HK$）：",
            packaging_subtotal_row,
            lift_cost_row - 1,
        )
        packaging_fx_row = find_between(
            3,
            "美金兑港币汇率：",
            packaging_subtotal_row,
            lift_cost_row - 1,
        )
        packaging_usd_row = find_between(
            3,
            "包装成本（USD）：",
            packaging_subtotal_row,
            lift_cost_row - 1,
        )
        assert sheet.cell(packaging_title_row, 5).value == "港柜"
        assert sheet.cell(packaging_title_row, 6).value == "港散"
        assert sheet.cell(packaging_freight_row, 5).value == 2.645
        assert sheet.cell(packaging_freight_row, 6).value == 5.447
        assert sheet.cell(packaging_subtotal_row, 5).value == (
            f"=$D${packaging_subtotal_row}+E{packaging_freight_row}"
        )
        assert sheet.cell(packaging_subtotal_row + 1, 4).value == 1.25
        assert sheet.cell(packaging_hkd_row, 5).value == (
            f"=E{packaging_subtotal_row}*E{packaging_subtotal_row + 1}/E{packaging_subtotal_row + 2}"
        )
        assert sheet.cell(packaging_fx_row, 5).value == "=$R$4"
        assert sheet.cell(packaging_usd_row, 5).value == (
            f"=E{packaging_hkd_row}/E{packaging_fx_row}"
        )
        assert sheet.cell(lift_cost_row, 5).value == 3.657
        assert sheet.cell(lift_cost_row, 6).value == 10.169
        lift_hkd_row = lift_cost_row + 3
        lift_fx_row = lift_cost_row + 4
        lift_usd_row = lift_cost_row + 5
        assert sheet.cell(lift_hkd_row, 5).value == (
            f"=E{lift_cost_row}*E{lift_cost_row + 1}/E{lift_cost_row + 2}"
        )
        assert sheet.cell(lift_fx_row, 5).value == "=$R$4"
        assert sheet.cell(lift_usd_row, 5).value == f"=E{lift_hkd_row}/E{lift_fx_row}"
        assert not any(
            sheet.cell(row, 3).value in {
                "分段汇总及运输方案（HKD）",
                "包装、纸箱及运输（一次汇总）",
                "分段报价合计（HKD）",
                "综合报价（USD，最后换算）",
            }
            for row in range(1, sheet.max_row + 1)
        )
        freight_total_row = _find_row(sheet, 3, "产品价（含运费）（USD）：")
        combined_total_row = _find_row(sheet, 3, "产品价（含吊柜费）（USD）：")
        final_header_row = freight_total_row - 1
        assert lift_usd_row < final_header_row < freight_total_row < combined_total_row
        assert sheet.cell(final_header_row, 4).value == "出厂价"
        assert sheet.cell(final_header_row, 5).value == "港柜"
        assert sheet.cell(freight_total_row, 5).value == (
            f"=$D${main_usd_row}+$D${mirror_usd_row}+E{packaging_usd_row}"
        )
        assert sheet.cell(combined_total_row, 4).value == f"=D{freight_total_row}"
        assert sheet.cell(combined_total_row, 5).value == (
            f"=E{freight_total_row}+E{lift_usd_row}"
        )
        legacy_header_row = _find_row(sheet, 3, "旺季价")
        assert sheet.cell(legacy_header_row + 1, 4).value == (
            f"=F{combined_total_row}*$R$4"
        )
        assert str(sheet.cell(legacy_header_row + 1, 5).value).count("SUMIF") >= 3
        third_header_row = legacy_header_row + 6
        assert str(sheet.cell(third_header_row + 1, 8).value).count("SUMIF") >= 3
        tax_amount_row = legacy_header_row + 10
        assert str(sheet.cell(tax_amount_row, 4).value).count("SUMIF") >= 3
        assert "_报价明细计算" not in str(sheet.cell(tax_amount_row, 4).value)
        # Saving the deliverable must preserve real, visible Excel formulas, not text/results.
        with BytesIO() as output:
            workbook.save(output)
            output.seek(0)
            saved = load_workbook(output, data_only=False)
            try:
                for detail_row in (adhesive_row, paper_pallet_row):
                    cell = saved["报价明细"].cell(detail_row, 4)
                    assert cell.data_type == "f"
                    assert cell.value == sheet.cell(detail_row, 4).value
                    assert "_报价明细计算" not in cell.value
            finally:
                saved.close()
    finally:
        workbook.close()
