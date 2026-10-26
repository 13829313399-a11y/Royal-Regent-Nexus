"""Order reference quantities and editable receipt facts stay separate in XLSX templates."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from fastapi import HTTPException
from sqlalchemy.orm import Session

from test_indonesia_materials import (client, actor, command, material, product, order, supply,
                                     allocation, historical_allocation, shipment, workspace, BASE)
from test_indonesia_warehouse_import import choose, confirm, source
from app.models.indonesia_materials import IndonesiaMaterialRecord as Record
from app.services.indonesia_warehouse_import import parse_warehouse_receipts
from app.services.indonesia_receipt_template import HEADERS, HEADER_ROW, FIRST_ROW, SHEET, TITLE


def catalog(client, **changes):
    m = material(client, **changes)
    product(client, m, code='00040149', daily_capacity='1000000000000')
    return m, order(client, item_no='00040149', po_no='0000123')


def body(*orders, **changes):
    return dict(factory_id='huaxing', customer='合成客户', recipient='RRI', warehouse='胶件仓',
                selections=[dict(order_id=o['id'], demand_id=d['id']) for o in orders for d in o['data']['demands']]) | changes


def download(client, *orders, status=200, **changes):
    response = client.post(BASE + '/warehouse-imports/template', json=body(*orders, **changes))
    assert response.status_code == status, response.text if response.status_code != 200 else ''
    return response


def values(content, row=FIRST_ROW):
    book = openpyxl.load_workbook(BytesIO(content))
    return {name: book[SHEET].cell(row, index).value for index, name in enumerate(HEADERS, 1)}


def fill(content, changes, row=FIRST_ROW):
    book = openpyxl.load_workbook(BytesIO(content))
    sheet = book[SHEET]
    for key, value in changes.items():
        sheet.cell(row, HEADERS.index(key) + 1, value)
    stream = BytesIO(); book.save(stream)
    return stream.getvalue()


def test_template_has_official_names_reference_amounts_and_blank_editable_receipts(client):
    m, o = catalog(client)
    response = download(client, o)
    assert response.headers['content-type'].startswith('application/vnd.openxmlformats')
    assert response.headers['cache-control'] == 'no-store'
    data = values(response.content)
    assert data['物料编码'] == '00001' and data['货号'] == '00040149' and data['客户PO'] == '0000123'
    assert data['正式报关中文名称'] == m['data']['name'] and data['正式报关英文名称'] == m['data']['english_name']
    assert data['规格'] == 'TEST' and data['数量单位'] == 'PCS'
    assert data['本次订单物料需求数量'] == data['剩余未送货数量'] == '100'
    assert data['已实收并分配数量'] == '0'
    assert data['仓库'] == '胶件仓'
    assert all(data[key] is None for key in ('收料单号', '行号', '实收日期', '正常实收数量', '备品数量', '箱号', '箱数'))
    assert not any('别名' in key or '原名称' in key or '价格' in key or '单价' in key for key in data)
    sheet = openpyxl.load_workbook(BytesIO(response.content))[SHEET]
    assert sheet['A1'].value == TITLE and sheet.freeze_panes == 'A6'
    assert sheet['M6'].fill.fgColor.rgb.endswith('F0F3F6') and sheet['T6'].fill.fgColor.rgb.endswith('FFF4CC')
    assert sheet['M6'].protection.locked and not sheet['T6'].protection.locked
    assert sheet['H6'].number_format == '@' and sheet.auto_filter.ref == 'A5:AA6'
    assert sheet['I6'].border.bottom.style == 'thin'
    assert all(sheet.column_dimensions[column].hidden for column in ('D', 'E', 'G'))
    target = Path('../.local-dev/receipt-template-example.xlsx')
    target.parent.mkdir(exist_ok=True, parents=True); target.write_bytes(response.content)


def test_reference_coverage_counts_received_normal_and_spares_only_and_never_readds_shipped(client):
    m, o = catalog(client)
    other = order(client, contract='TEST-B', item_no='00040149', po_no='PO-OTHER')
    received = command(client, 'receipt_confirm', dict(actual_date='2026-10-22'), supply(client, m))
    ours = allocation(client, received, o, normal='30', spare='5')
    allocation(client, received, other, normal='40', spare='0')
    expected = supply(client, m, normal='10', spare='0', batch='EXPECTED', purchase_no='P-EXPECTED', purchase_quantity='10')
    historical_allocation(client, expected, o, normal='10', spare='0')
    draft = shipment(client, ours)
    before = values(download(client, o).content)
    command(client, 'shipment_confirm', {}, draft)
    after = values(download(client, o).content)
    assert before['已实收并分配数量'] == after['已实收并分配数量'] == '35'
    assert before['剩余未送货数量'] == after['剩余未送货数量'] == '65'
    board = next(item for item in workspace(client)['board'] if item['id'] == o['id'])
    assert after['剩余未送货数量'] == board['demands'][0]['shortage']


def test_reference_uses_final_adjusted_demand_and_clamps_legacy_overcoverage(client):
    m, o = catalog(client)
    o = command(client, 'demand_adjust', dict(demand_id=o['data']['demands'][0]['id'], quantity='75', reason='实际订单调整'), o)
    sup = command(client, 'receipt_confirm', dict(actual_date='2026-10-22'), supply(client, m))
    historical_allocation(client, sup, o)
    data = values(download(client, o).content)
    assert data['本次订单物料需求数量'] == '75'
    assert data['已实收并分配数量'] == '100' and data['剩余未送货数量'] == '0'


def test_formula_looking_names_are_literal_and_eight_decimal_quantities_keep_precision(client):
    _, first = catalog(client, name='=危险公式', english_name='+SUM(1,2)')
    second = order(client, item_no='00040149', contract='TEST-LARGE', quantity='123456789012.12345678')
    sheet = openpyxl.load_workbook(BytesIO(download(client, second).content))[SHEET]
    assert sheet['I6'].value == '=危险公式' and sheet['I6'].data_type == 's'
    assert sheet['J6'].value == '+SUM(1,2)' and sheet['J6'].data_type == 's'
    assert sheet['M6'].value == sheet['O6'].value == '123456789012.12345678'


def test_same_material_in_two_contracts_stays_two_rows_and_unfilled_row_is_skipped(client):
    _, first = catalog(client)
    second = order(client, contract='TEST-B', item_no='00040149', po_no='PO-OTHER')
    content = download(client, first, second).content
    assert values(content)['合同号'] == 'TEST-A' and values(content, FIRST_ROW + 1)['合同号'] == 'TEST-B'
    content = fill(content, {'收料单号': 'WH-001', '行号': '0001', '实收日期': '2026-10-26', '正常实收数量': '90', '备品数量': '10'})
    rows = parse_warehouse_receipts(content, '回传.xlsx')
    assert len(rows) == 1 and not rows[0]['source_issues']
    assert rows[0]['source']['normal_quantity'] == '90' and rows[0]['source']['spare_quantity'] == '10'
    assert rows[0]['source']['receipt_line'] == '0001' and rows[0]['source']['po_no'] == '0000123'
    assert '剩余未送货数量' not in rows[0]['source']
    response = client.post(BASE + '/warehouse-imports/preview', data=dict(factory_id='huaxing', customer='合成客户', recipient='RRI', warehouse=''),
                           files={'file': ('回传.xlsx', content, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code == 200, response.text
    preview = response.json()['record']
    row = preview['data']['rows'][0]
    assert row['auto_material_id'] and len(row['targets'][row['auto_material_id']]) == 1
    purchase = {key: source()[key] for key in ('purchase_no', 'purchase_line', 'purchase_date', 'purchase_quantity', 'supplier', 'price', 'currency')}
    selection = choose(preview, purchase=purchase)
    confirmed = confirm(client, preview, [selection])
    assert len(confirmed['data']['shipment_ids']) == 1
    assert next(o for o in workspace(client)['board'] if o['id'] == first['id'])['demands'][0]['coverage'] == '100'


def test_blank_template_is_not_receipt_and_partial_entry_without_quantity_reports_errors(client):
    _, o = catalog(client)
    content = download(client, o).content
    with pytest.raises(HTTPException, match='尚未填写本次实收数量'):
        parse_warehouse_receipts(content, '空模板.xlsx')
    for change in ({'收料单号': 'WH-01'}, {'箱数': '1'}, {'正常实收数量': '0'}, {'备品数量': '3'}):
        rows = parse_warehouse_receipts(fill(content, change), '半填写.xlsx')
        assert len(rows) == 1 and rows[0]['source_issues']
    # Old-format documents still surface missing facts, rather than silently disappearing.
    book = openpyxl.Workbook(); sheet = book.active
    sheet.append(['物料名称', '物料编码', '正常实收数量']); sheet.append(['合成螺丝', '00001', None])
    stream = BytesIO(); book.save(stream)
    rows = parse_warehouse_receipts(stream.getvalue(), '旧表.xlsx')
    assert len(rows) == 1 and rows[0]['source_issues']


@pytest.mark.parametrize('changes,status', [({'customer': '另一客人'}, 422), ({'recipient': 'RRM'}, 422),
    ({'factory_id': 'huadeng'}, 403), ({'factory_id': 'group'}, 422), ({'selections': []}, 422)])
def test_template_rejects_mismatched_scope_and_empty_selection(client, changes, status):
    _, o = catalog(client)
    download(client, o, status=status, **changes)


def test_template_rejects_cross_factory_record_duplicate_missing_deleted_and_oversized_selection(client):
    m, o = catalog(client)
    selected = body(o)['selections']
    download(client, o, selections=selected * 2, status=422)
    download(client, o, selections=selected * 201, status=422)
    download(client, o, selections=[dict(order_id=o['id'], demand_id='missing')], status=409)
    client.materials_auth['user'] = actor('read', 'write', factory='huadeng')
    download(client, o, factory_id='huadeng', status=404)
    client.materials_auth['user'] = actor('read', 'write')
    with Session(client.engine) as db:
        record = db.get(Record, m['id']); record.data = record.data | {'deleted': True}; db.commit()
    download(client, o, status=404)


@pytest.mark.parametrize('actions,department', [(('read',), 'sales-business'), (('read', 'receipt_confirm'), 'pmc-warehouse')])
def test_template_requires_existing_sales_write_permission(client, actions, department):
    _, o = catalog(client)
    client.materials_auth['user'] = actor(*actions, department=department)
    download(client, o, status=403)


def test_legacy_csv_template_stays_available(client):
    response = client.get(BASE + '/warehouse-imports/template', params={'factory_id': 'huaxing'})
    assert response.status_code == 200 and response.content.startswith(b'\xef\xbb\xbf')
    assert '正常实收数量' in response.content.decode('utf-8-sig')


@pytest.mark.parametrize('context,message', [({'customer': '另一客人'}, '原单客户'), ({'recipient': 'RRM'}, '原单接收方')])
def test_uploaded_template_context_cannot_be_redirected_to_identical_contract_in_another_context(client, context, message):
    m, o = catalog(client)
    content = fill(download(client, o).content, {'收料单号': 'WH-001', '行号': '1', '实收日期': '2026-10-26', '正常实收数量': '90', '备品数量': '10'})
    if context.get('customer'):
        product(client, m, code='00040149', customer=context['customer'])
    order(client, item_no='00040149', po_no='0000123', **context)
    response = client.post(BASE + '/warehouse-imports/preview', data=dict(factory_id='huaxing', customer='合成客户', recipient='RRI', warehouse='') | context,
                           files={'file': ('回传.xlsx', content, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code == 200, response.text
    row = response.json()['record']['data']['rows'][0]
    assert row['blocked'] and any(message in issue for issue in row['issues'])
