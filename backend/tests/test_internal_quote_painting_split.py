from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest

from app.services.internal_quote_import import parse_internal_quote_workbook
from app.services.internal_quote_calculator import calculate_section, CalculationInputError
from app.services.internal_quote_templates import build_internal_quote_import_template


def source(paint=2, labor='=E2-F2', *, headers=None):
    w = openpyxl.Workbook()
    w.active.append(headers or ['名称', '位置', 'UV', 'UV单价', '总报价', '油漆', '人工', '备注'])
    w.active.append(['外壳', '正面', 2, 5, '=C2*D2', paint, labor, '保持原备注'])
    out = BytesIO()
    w.save(out)
    return out.getvalue()


def calculate(rows):
    return calculate_section('painting', {'rows': rows}, {}, 'TEST')


def test_new_template_and_uncached_formula_map_split_and_recalculate_after_edit():
    data, _ = build_internal_quote_import_template('painting')
    w = openpyxl.load_workbook(BytesIO(data))
    assert [w.active.cell(2, c).value for c in range(22, 26)] == ['总报价', '油漆', '人工', '备注']
    assert w.active['X3'].value == '=V3-W3'
    assert len(w.active._images) == 14
    real_rows = parse_internal_quote_workbook(data, 'painting').payload_fragment['rows']
    assert not any(row['source_row'] == 41 for row in real_rows)
    first = calculate([real_rows[0]])['totals']
    assert first == {'total_hkd': '0.3000', 'painting_labor_hkd': '0.2800', 'paint_material_hkd': '0.0200'}
    row = parse_internal_quote_workbook(source(), 'painting').payload_fragment['rows'][0]
    assert row['cost_allocation'] == 'split'
    assert row['paint_cost_hkd'] == '2'
    assert row['labor_cost_hkd'] is None
    assert row['remark'] == '保持原备注'
    row.update(pricing_component_id='shell', markup_override='1.4')
    calc = calculate([row])
    assert calc['totals'] == {'total_hkd': '10.0000', 'painting_labor_hkd': '8.0000', 'paint_material_hkd': '2.0000'}
    assert [line['kind'] for line in calc['line_breakdown']] == ['painting_quick_labor', 'painting_quick_paint']
    assert all(line['pricing_component_id'] == 'shell' and line['markup_override'] == '1.4000' for line in calc['line_breakdown'])
    row['operations']['uv']['quantity'] = '3'
    assert calculate([row])['totals']['painting_labor_hkd'] == '13.0000'


@pytest.mark.parametrize('paint,labor,expected', [(0, None, ('0.0000', '10.0000')), (None, 3, ('7.0000', '3.0000')), (2, 8, ('2.0000', '8.0000'))])
def test_explicit_zero_and_single_value_are_preserved(paint, labor, expected):
    row = parse_internal_quote_workbook(source(paint, labor), 'painting').payload_fragment['rows'][0]
    totals = calculate([row])['totals']
    assert (totals['paint_material_hkd'], totals['painting_labor_hkd']) == expected


@pytest.mark.parametrize('paint,labor', [(-1, 11), ('NaN', 8), ('待报价2', 8), (2, 9), (11, None), (2, '=UNKNOWN(E2)')])
def test_bad_split_fails_instead_of_silent_zero(paint, labor):
    with pytest.raises(ValueError):
        parse_internal_quote_workbook(source(paint, labor), 'painting')


def test_blank_pair_is_saved_as_pending_not_legacy_allocation():
    row = parse_internal_quote_workbook(source(None, None), 'painting').payload_fragment['rows'][0]
    assert row['paint_cost_hkd'] is None and row['labor_cost_hkd'] is None
    calc = calculate([row])
    assert calc['status'] == 'blocked'
    assert 'painting_split_missing' in {item['code'] for item in calc['warnings']}
    row.update(paint_cost_hkd=2, labor_cost_hkd=9)
    with pytest.raises(CalculationInputError, match='合计'):
        calculate([row])


def test_legacy_and_direct_history_mix_without_new_tax_or_duplicate_cost():
    new = parse_internal_quote_workbook(source(), 'painting').payload_fragment['rows'][0]
    legacy = {'operations': {'uv': {'quantity': 1, 'unit_price_hkd': 10}}}
    direct = {'cost_allocation': 'direct', 'paint_hkd': 100, 'spray_labor_hkd': 100,
              'operations': {'paint': {'quantity': 1, 'unit_price_hkd': '1.13'}}}
    totals = calculate([new, legacy, direct])['totals']
    assert totals == {'total_hkd': '21.1300', 'painting_labor_hkd': '15.0000', 'paint_material_hkd': '6.1300'}
    assert Decimal(totals['total_hkd']) == Decimal(totals['paint_material_hkd']) + Decimal(totals['painting_labor_hkd'])


def test_reordered_headers_do_not_mix_remarks_and_prices():
    w = openpyxl.Workbook()
    w.active.append(['名称', '位置', '油漆 HKD', '备注', '人工 HKD', 'UV单价', 'UV'])
    w.active.append(['壳', '头', 2, '说明', 8, 5, 2])
    out = BytesIO(); w.save(out)
    row = parse_internal_quote_workbook(out.getvalue(), 'painting').payload_fragment['rows'][0]
    assert (row['paint_cost_hkd'], row['labor_cost_hkd'], row['remark']) == ('2', '8', '说明')


@pytest.mark.parametrize('mode', ['ordinary', 'component', 'detached'])
def test_final_summary_excel_classification_and_markups(mode):
    from types import SimpleNamespace
    from app.services.internal_quote import _rr2_cost_summary
    from app.services.internal_quote_excel import _build_summary_sheet
    from test_internal_quote_excel_formula_links import _section
    rows = parse_internal_quote_workbook(source(), 'painting').payload_fragment['rows']
    rows[0]['pricing_component_id'] = 'a'
    if mode == 'detached':
        rows[0]['markup_override'] = '1.4'
    rows.append({**rows[0], 'name': '副壳', 'position': '背面', 'pricing_component_id': 'b',
                 'paint_cost_hkd': '4', 'labor_cost_hkd': '6'})
    payload = {'rows': rows}
    calc = calculate(rows)
    sales = {'pricing_mode': 'component', 'pricing_components': [
        {'id': 'a', 'name': '主体', 'markup_x': '1.2'}, {'id': 'b', 'name': '配件', 'markup_x': '1.4'}]} if mode == 'component' else {}
    sections = [_section('painting', payload=payload, calculation=calc), _section('sales', payload=sales)]
    for section in sections:
        section.is_required = True
    snapshot = {'fx': {'rmb_hkd': '.85', 'hkd_usd': '7.8'}, 'markup': '1.2', 'settlement': '1'}
    context = {'painting_hkd': Decimal('20'), 'factory_price_hkd': Decimal('20')}
    summary = _rr2_cost_summary(sections, context, snapshot, factory_id='huakang-b' if mode == 'component' else 'huakang-a')
    costs = {item['key']: item['value'] for item in summary['t3']}
    assert (costs['paint_material'], costs['painting_labor']) == ('6.0000', '14.0000')
    pricing = summary['shipping_pricing']
    assert sum(Decimal(e['amount_hkd']) for e in pricing['pricing_entries']) == 20
    assert sum(Decimal(g['quoted_hkd']) for g in pricing['pricing_groups']) == {'ordinary': 24, 'component': 26, 'detached': 28}[mode]
    w = openpyxl.Workbook()
    _build_summary_sheet(w, SimpleNamespace(product_name='喷油拆分验证', quote_no='PAINT-SPLIT', customer='JustPlay' if mode == 'component' else '普通客', region_code='mainland', remark='测试'), sections, snapshot, summary, context, [])
    s = w['报价明细']
    # Both ordinary and component outputs classify raw costs in B:D.
    import re
    values = {category: sum(Decimal(str(s.cell(r, 4).value).removeprefix('=')) for r in range(1, s.max_row + 1)
                            if s.cell(r, 2).value == category and re.fullmatch(r'=?\d+(?:\.\d+)?', str(s.cell(r, 4).value)))
              for category in ['油漆', '喷油工']}
    assert values == {'油漆': Decimal('6'), '喷油工': Decimal('14')}
