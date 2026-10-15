import json
from datetime import datetime
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.services import internal_quote_excel as excel


@pytest.mark.parametrize("customer,factory,release,expected", [
    ("BuzzBee", "huaxing", "p4_final_approved", True),
    ("BuzzBee", "huaxing", "p4_direct_issued", True),
    ("BuzzBee", "huaxing", "p3", False),
    ("BuzzBee", "huakang-a", "p4_final_approved", False),
    ("银辉", "huaxing", "p4_final_approved", False),
])
def test_buzzbee_handoff_uses_creation_identity_only(monkeypatch, customer, factory, release, expected):
    quote = SimpleNamespace(quote_no="IQ-BB", customer=customer, factory_id=factory,
                            product_name="本方案产品", version_label="R2", formula_version="F1",
                            reference_snapshot_id="REF1", header_revision=3,
                            created_at=datetime(2026, 9, 2, 10, 20))
    # Layout is covered separately; exercise the actual export gate and structured writer.
    monkeypatch.setattr(excel, "_build_summary_sheet", lambda *args: None)
    for name in ("_build_electronic_sheet", "_build_sewing_sheet", "_build_hair_sheet", "_build_assembly_sheet"):
        monkeypatch.setattr(excel, name, lambda *args: None)
    data = excel.build_internal_quote_workbook(quote, [], {"release_stage": release}, {})
    book = load_workbook(BytesIO(data))
    rows = list(book["结构化数据"].values) if "结构化数据" in book.sheetnames else []
    chunks = [row[10] for row in rows if row[0] == "customer_mapping"]
    assert bool(chunks) is expected
    if expected:
        mapping = json.loads("".join(chunks))
        assert mapping == {"version": "buzzbee-v2", "factory_id": "huaxing", "quote_no": "IQ-BB",
                           "version_label": "R2", "customer": "BuzzBee", "product_name": "本方案产品",
                           "quote_date": "2026-09-02", "formula_version": "F1", "reference_snapshot_id": "REF1"}
        assert book["结构化数据"].sheet_state == "veryHidden"


def test_structured_handoff_retains_split_inputs_and_legacy_customer_fields():
    quote = SimpleNamespace(header_revision=3, reference_snapshot_id="REF1")
    payloads = {
        "engineering": {"molds": [{
            "id": "mold-1", "item": "上壳/下壳", "mold_specification": "历史模具规格",
            "material": "历史材质", "cycle_time_seconds": 20,
            "parts": [
                {"id": "upper", "name": "上壳", "unit_net_weight_g": 12, "output_count": 4, "quantity": 2},
                {"id": "lower", "name": "下壳", "unit_net_weight_g": 10, "output_count": 2, "quantity": 1},
            ],
        }]},
        "molding": {"injection_lines": [
            {"item": "上壳", "net_weight_g": 12, "cavity": 4, "sets": 2, "quantity": 3, "cycle_time_seconds": 20},
            {"item": "下壳", "net_weight_g": 10, "cavity": 2, "sets": 1, "quantity": 3, "cycle_time_seconds": 20},
        ]},
    }
    sections = [SimpleNamespace(
        department=code, department_name=code, status="draft", revision=2,
        calculation_status="valid", dependency_status="current", calculation_hash=f"hash-{code}",
        is_required=True, payload_json=json.dumps(payload, ensure_ascii=False),
        calculation_json=json.dumps({"total_hkd": 1.25}),
    ) for code, payload in payloads.items()]
    book = Workbook()
    excel._build_structured_data_sheet(book, quote, sections, {"fx": {"rmb_to_hkd": 0.85}})
    buffer = BytesIO()
    book.save(buffer)
    rows = list(load_workbook(BytesIO(buffer.getvalue()))["结构化数据"].values)

    for code, payload in payloads.items():
        records = [row for row in rows if row[0] == "payload" and row[1] == code]
        assert json.loads("".join(row[10] for row in records)) == payload
        assert all(row[3:8] == ("draft", 2, "valid", "current", f"hash-{code}") for row in records)
        calculations = [row[10] for row in rows if row[0] == "calculation" and row[1] == code]
        assert json.loads("".join(calculations)) == {"total_hkd": 1.25}
