"""Purchase views must reuse frozen special versions without changing history."""
from io import BytesIO
from openpyxl import load_workbook
import pytest
from test_carton_special_requests import make_client, setup, ok, create_body, command, login_as, supplier_login, accept, BASE, SUP, FACTORY


@pytest.fixture(autouse=True)
def isolated_database_directory(monkeypatch, tmp_path):
    import test_molding_sample_api
    monkeypatch.setattr(test_molding_sample_api, "TEST_TMP_DIR", tmp_path)


@pytest.mark.parametrize("kind", ["SAMPLE", "WORKSHOP_MISC"])
def test_frozen_versions_shared_export_supplier_access_and_history(monkeypatch, kind):
    with make_client(monkeypatch) as client:
        setup(client)
        body = create_body(kind)
        body["requester_name"] = ""
        row = ok(client.post(BASE, json=body), 201)
        url = BASE + f"/{row['id']}/purchase-orders"
        assert ok(client.get(url, params={"factory_id": FACTORY}))["issues"] == []
        row = ok(client.post(BASE + f"/{row['id']}/send", json=command(row)))
        initial = ok(client.get(url, params={"factory_id": FACTORY}))["issues"][0]
        export_url = url + f"/{initial['id']}/export"
        original = client.get(export_url, params={"factory_id": FACTORY})
        assert original.status_code == 200, original.text
        sheet = load_workbook(BytesIO(original.content)).active
        assert sheet["G10"].value == 100 and sheet["H10"].value == 100
        assert sheet["F5"].value == ("样板单" if kind == "SAMPLE" else "杂单")
        before_history = row["history"]
        edit = create_body(kind) | command(row) | {"reason": "增加实际需求", "lines": [dict(row["lines"][0], quantity="120")]}
        # Response-only fields are not command input.
        edit["lines"] = [{k: v for k, v in line.items() if k in {"id", "item_no", "packaging_type", "paper_quality", "specification", "dimension_unit", "quantity", "unit", "mark_mode"}} for line in edit["lines"]]
        row = ok(client.post(BASE + f"/{row['id']}/adjust", json=edit))
        assert row["history"][:len(before_history)] == before_history
        context = ok(client.get(url, params={"factory_id": FACTORY}))
        assert len(context["issues"]) == 2
        assert context["issues"][0]["document_type"] == "APPEND"
        latest = client.get(url + f"/{context['issues'][0]['id']}/export", params={"factory_id": FACTORY})
        sheet = load_workbook(BytesIO(latest.content)).active
        assert [sheet[f"{col}10"].value for col in "FGH"] == [100, 20, 120]
        old = load_workbook(BytesIO(client.get(export_url, params={"factory_id": FACTORY}).content)).active
        assert [old[f"{col}10"].value for col in "FGH"] == [0, 100, 100]
        # Internal and supplier views use the same immutable document identity.
        supplier_login(client)
        external_url = SUP + f"/{row['id']}/purchase-orders"
        assert ok(client.get(external_url, params={"factory_id": FACTORY}))["issues"] == context["issues"]
        assert client.get(external_url, params={"factory_id": "huakang-b"}).status_code in (403, 404)
        assert client.get(external_url + '/missing/export', params={"factory_id": FACTORY}).status_code == 404
        accept(client, row)
        docs = ok(client.get('/api/carton-supplier/documents', params={"factory_id": FACTORY}))
        owned = [doc for doc in docs if doc.get("special_request_id") == row["id"]]
        assert {doc["id"] for doc in owned} == {item["id"] for item in context["issues"]}
        assert owned[0]["lines"][0]["change_quantity"] == "20.0000"
        assert owned[0]["supplier_acceptance"]["status"] == "ACCEPTED"
        assert owned[0]["supplier_acceptance"]["accepted_line_count"] == 1
        assert owned[1]["supplier_acceptance"]["status"] == "HISTORICAL"
        selection = {"documents": [{"factory_id": FACTORY, "kind": "PURCHASE", "id": owned[0]["id"]}]}
        exported = client.post('/api/carton-supplier/documents/export.xlsx', json=selection)
        assert exported.status_code == 200, exported.text
        assert client.post('/api/carton-supplier/documents/order-import.xlsx', json=selection).status_code == 422
        after_export = ok(client.get('/api/carton-supplier/documents', params={"factory_id": FACTORY}))
        assert next(doc for doc in after_export if doc["id"] == owned[0]["id"])["export_count"] == 1
        login_as(client, "admin")
        unchanged = ok(client.get(BASE + f"/{row['id']}", params={"factory_id": FACTORY}))
        assert unchanged["revision"] == row["revision"] and unchanged["history"] == row["history"]
        assert unchanged["lines"][0]["id"] == row["lines"][0]["id"]
        cancelled = ok(client.post(BASE + f"/{row['id']}/cancel", json=command(row)))
        history = ok(client.get(url, params={"factory_id": FACTORY}))["issues"]
        assert len(history) == 3 and history[0]["is_cancellation"]
        assert cancelled["history"][:len(row["history"])] == row["history"]
        # Neither another factory nor an unrelated own-request-only user may export.
        from test_molding_sample_api import ensure_test_user
        from test_carton_supplier_portal import grant_supplier_permissions
        ensure_test_user("engineer_peer")
        grant_supplier_permissions("user-engineer-peer", permissions=("carton_procurement:special_request_submit",))
        login_as(client, "engineer_peer")
        assert client.get(url, params={"factory_id": FACTORY}).status_code == 403
        assert client.get(export_url, params={"factory_id": FACTORY}).status_code == 403
