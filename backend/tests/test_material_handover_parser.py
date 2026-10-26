import importlib
import re
from io import BytesIO
from decimal import Decimal
from zipfile import ZIP_DEFLATED, ZipFile
import pytest
from test_material_handover import workbook


def test_recognition_suggests_editable_categories_without_changing_original_units():
    parser = importlib.import_module('app.services.material_handover_parser')
    units = ['码', '米', 'Pcs', '卷', '千克(转算专用)', '片', ' PＣＳ ']
    result = parser.parse('分类建议.xlsx', workbook(rows=[
        [str(index), '同名物料', '2', unit, ''] for index, unit in enumerate(units, 1)]))
    lines = result['data']['lines']
    assert [line['material_category'] for line in lines] == ['FABRIC', 'FABRIC', 'ACCESSORY', 'ACCESSORY', '', '', 'ACCESSORY']
    assert [line['unit'] for line in lines] == [unit.strip() for unit in units]
    assert [line['quantity'] for line in lines] == ['2'] * len(units)


def test_merged_delivery_footers_are_not_materials_and_keep_original_evidence():
    from openpyxl import Workbook
    parser = importlib.import_module('app.services.material_handover_parser')
    book = Workbook(); sheet = book.active; sheet.title = '销售出货单'
    sheet.append(['订单 / 货号', '布料名称-色号', '1', '2', '合计数量'])
    quantities = [2000, 152, 448, 2000, 830, 2000, 490]
    for index, quantity in enumerate(quantities):
        sheet.append([f'订单1 缸号：{index // 3 + 1}#', '01-01-10-00002 黑布', quantity, None, f'{quantity}Y'])
    sheet.append(['合计数量', None, None, None, sum(quantities)])
    sheet.merge_cells('A9:D9')
    footers = ['备注：请核对布种、色号、缸号及各卷数量。', '送货人：____________', '测试样本：资料为模拟数据，仅用于测试。']
    for text in footers:
        sheet.append([text]); sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=5)
    output = BytesIO(); book.save(output)
    result = parser.parse('供应商原单.xlsx', output.getvalue())
    lines = result['data']['lines']
    assert [line['quantity'] for line in lines] == [str(value) for value in quantities]
    assert len(lines) == 7
    assert sum(Decimal(line['quantity']) for line in lines) == Decimal('7920')
    assert [line['source_line'] for line in lines] == [f'销售出货单:{index}' for index in range(2, 9)]
    assert set(result['evidence']) == {line['id'] for line in lines}
    assert all(text in result['raw_text'] for text in ['合计数量', *footers])
    assert not any('缺少' in warning for warning in result['warnings'])


@pytest.mark.parametrize('label', ['合计', '合计数量', '数量合计', '本页小计', '总计数量', 'Grand Total', 'Subtotal', 'Total Quantity'])
@pytest.mark.parametrize('identity_field', ['item_code', 'item_name'])
def test_summary_labels_in_either_identity_column_are_not_materials(label, identity_field):
    parser = importlib.import_module('app.services.material_handover_parser')
    footer = [label, '', '40', '码', ''] if identity_field == 'item_code' else ['', label, '40', '码', '']
    result = parser.parse('自由布局.xlsx', workbook(rows=[['A', '布料', '40', '码', ''], footer]))
    assert [(line['item_code'], line['quantity']) for line in result['data']['lines']] == [('A', '40')]


def test_footer_filter_keeps_incomplete_materials_names_and_later_sections():
    parser = importlib.import_module('app.services.material_handover_parser')
    rows = [['A', '备注：特殊布料', '2', '码', ''], ['MAT-TOTAL', '合计数量', '3', '码', ''],
            ['', '备注布', '4', '米', ''], ['C', '正常物料', '', '', ''],
            ['', '小计', '9', '码', ''], ['', '审核人：张三', '', '', ''], ['D', '下一部分布料', '1', '码', '']]
    result = parser.parse('多段物料.xlsx', workbook(rows=rows))
    assert [line['item_code'] for line in result['data']['lines']] == ['A', 'MAT-TOTAL', '', 'C', 'D']
    assert result['data']['lines'][2]['item_name'] == '备注布'
    assert result['data']['lines'][3]['quantity'] == ''
    assert any('缺少' in warning for warning in result['warnings'])


def test_populated_standard_template_signature_area_is_not_material():
    from openpyxl import Workbook
    parser = importlib.import_module('app.services.material_handover_parser')
    book = Workbook(); sheet = book.active
    sheet.append(['物料分类', '物料编码', '物料名称', '规格', '颜色', '数量', '单位', '缸号', '卷号', '生产单号', '单价', '币种', '备注'])
    sheet.append(['布料', 'F001', '浅蓝色布', '', '', 100.5, '码'])
    sheet.append(['辅料', 'A001', '开单员：特殊辅料', '', '', 2, 'Pcs'])
    for label_cell, label_range, label, value_cell, value in [
        ('A5', 'A5:C5', '开单员：', 'D5', '张三'),
        ('F5', 'F5:G5', '送货人：', 'H5', '李四'),
        ('J5', None, '收货人：', 'K5', '王五'),
        ('A7', 'A7:C7', '签收日期：', 'D7', '2026-10-26'),
        ('F7', 'F7:G7', '签收备注：', 'H7', '外包装完好'),
    ]:
        sheet[label_cell] = label
        if label_range:
            sheet.merge_cells(label_range)
        sheet[value_cell] = value
    output = BytesIO(); book.save(output)
    result = parser.parse('统一送货单.xlsx', output.getvalue())
    assert [(line['item_code'], line['quantity']) for line in result['data']['lines']] == [('F001', '100.5'), ('A001', '2')]
    assert all(text in result['raw_text'] for text in ['开单员：', '送货人：', '收货人：', '签收日期：', '签收备注：', '张三', '李四', '王五', '2026-10-26', '外包装完好'])
    assert len(result['evidence']) == 2


@pytest.mark.parametrize('label', ['开单员', '開單員', '签收日期', '簽收日期', '签收备注', '簽收備註'])
def test_signature_labels_in_name_column_remain_original_evidence(label):
    parser = importlib.import_module('app.services.material_handover_parser')
    text = f'{label}：现场填写信息'
    result = parser.parse('签收原单.xlsx', workbook(rows=[['F001', '布料', '2', '码', ''], ['', text, '', '', '']]))
    assert [line['item_code'] for line in result['data']['lines']] == ['F001']
    assert text in result['raw_text']


def test_merged_material_and_dye_lots_keep_independent_rows_and_numbered_rolls():
    from openpyxl import Workbook
    parser = importlib.import_module('app.services.material_handover_parser')
    book = Workbook(); sheet = book.active
    sheet.append(['物料编码', '物料名称', '缸号', '1', '2', '合计数量'])
    sheet.append(['000123', '同种布料', '1#', 200, 200, '400Y'])
    sheet.append([None, None, None, 152, None, '152Y'])
    sheet.append([None, None, '2#', 200, 30, '230Y'])
    sheet.merge_cells('A2:A4'); sheet.merge_cells('B2:B4'); sheet.merge_cells('C2:C3')
    output = BytesIO(); book.save(output)
    result = parser.parse('合并多缸.xlsx', output.getvalue())
    lines = result['data']['lines']
    assert [(r['item_code'], r['lot'], r['quantity'], r['unit']) for r in lines] == [
        ('000123', '1#', '400', '码'), ('000123', '1#', '152', '码'), ('000123', '2#', '230', '码')]
    assert [r['roll_quantities'] for r in lines] == ['200 / 200', '152', '200 / 30']
    assert not any('逐卷数量' in warning for warning in result['warnings'])


def test_ocr_total_spanning_into_roll_header_is_not_counted_as_an_extra_roll():
    from app.services.document_tools.document_ir import DocumentIR, Table, Cell
    parser = importlib.import_module('app.services.material_handover_parser')
    table = Table(id='grid', row_count=2, column_count=6, cells=[
        Cell(id='h0',row=0,column=0,raw_text='物料编码'), Cell(id='h1',row=0,column=1,raw_text='物料名称'),
        Cell(id='h2',row=0,column=2,raw_text='1'), Cell(id='h3',row=0,column=3,raw_text='2'),
        Cell(id='h4',row=0,column=4,raw_text='3'), Cell(id='h5',row=0,column=5,raw_text='合计数量'),
        Cell(id='code',row=1,column=0,raw_text='000123'),Cell(id='name',row=1,column=1,raw_text='布料'),
        Cell(id='r1',row=1,column=2,raw_text='200'),Cell(id='r2',row=1,column=3,raw_text='30'),
        Cell(id='total',row=1,column=4,colspan=2,raw_text='230Y')])
    result = parser.map_ir(DocumentIR(source_type='image',tables=[table]))
    assert result['data']['lines'][0]['roll_quantities'] == '200 / 30'
    assert result['data']['lines'][0]['quantity'] == '230'


def test_mixed_units_first_row_and_planned_quantity_are_preserved():
    parser = importlib.import_module('app.services.material_handover_parser')
    result = parser.parse('budget.xlsx', workbook(headers=['产品名称','单位','物料编码','物料名称','单位','计划用量','总数量'],
        rows=[['产品','套','00012','胶水','支','40','40'], ['产品','套','00013','辅料','Pcs','73','73']], title='制单预算用料'))
    assert [(v['item_code'],v['unit'],v['quantity']) for v in result['data']['lines']] == [('00012','支','40'),('00013','Pcs','73')]
    assert any('预算' in warning for warning in result['warnings'])


def test_merged_ocr_headers_roll_quantities_and_mismatch_require_review():
    parser = importlib.import_module('app.services.material_handover_parser')
    ir = importlib.import_module('app.services.document_tools.document_ir')
    cells = [ir.Cell(id='head',row=0,column=0,rowspan=2,colspan=2,raw_text='订单号/货号\n布料名称-色号'),
        ir.Cell(id='rolls',row=0,column=2,colspan=2,raw_text='明细数量'),
        ir.Cell(id='total',row=0,column=4,rowspan=2,raw_text='合计\n数量'),
        ir.Cell(id='rh',row=1,column=2,colspan=2,raw_text='1\n2')]
    for column,value in enumerate(['订单1 缸号：2#','01-01-10-00002 黑布','200','201','400Y']):
        cells.append(ir.Cell(id=f'd{column}',row=2,column=column,raw_text=value))
    result = parser.map_ir(ir.DocumentIR(source_type='image', tables=[ir.Table(id='photo',row_count=3,column_count=5,cells=cells)]))
    assert len(result['data']['lines']) == 1
    row = result['data']['lines'][0]
    assert row['quantity'] == '400' and row['unit'] == '码' and row['lot'] == '2#'
    assert row['roll_quantities'] == '200 / 201'
    assert any('合计' in warning for warning in result['warnings'])


def test_docx_native_table_without_supplier_template():
    from docx import Document
    parser = importlib.import_module('app.services.material_handover_parser')
    doc = Document(); doc.add_paragraph('领料单 单号：REQ-002 2026-07-18')
    table = doc.add_table(rows=2,cols=4)
    for i,value in enumerate(['物料名称','单位','数量','物料编码']): table.cell(0,i).text=value
    for i,value in enumerate(['黑布','码','12.5','00012']): table.cell(1,i).text=value
    out=BytesIO(); doc.save(out)
    result=parser.parse('车间自由格式.docx',out.getvalue())
    assert result['data']['kind'] == 'ISSUE' and result['data']['lines'][0]['quantity'] == '12.5'


def test_disposable_worker_extracts_office_without_business_database():
    files = importlib.import_module('app.services.material_handover_files')
    result=files.extract('任意.xlsx', workbook())
    assert result['engine'] == 'xlsx' and result['data']['lines'][0]['quantity'] == '40'


def without_sheet_dimensions(source):
    output = BytesIO()
    with ZipFile(BytesIO(source)) as source_zip, ZipFile(output, 'w', ZIP_DEFLATED) as target_zip:
        for entry in source_zip.infolist():
            data = source_zip.read(entry.filename)
            if entry.filename.startswith('xl/worksheets/') and entry.filename.endswith('.xml'):
                data = re.sub(rb'<dimension[^>]*/>', b'', data)
            target_zip.writestr(entry, data)
    return output.getvalue()


def test_xlsx_without_sheet_dimension_preserves_delivery_details():
    parser = importlib.import_module('app.services.material_handover_parser')
    result = parser.parse('unsized.xlsx', without_sheet_dimensions(workbook()))
    assert result['data']['lines'][0]['item_code'] == '01-01-00001'
    assert result['data']['lines'][0]['quantity'] == '40'
    assert result['data']['source_no'] == 'DOC-001'


def test_xlsx_without_dimensions_can_include_empty_worksheets():
    from openpyxl import load_workbook
    parser = importlib.import_module('app.services.material_handover_parser')
    book = load_workbook(BytesIO(workbook()))
    book.create_sheet('空白页')
    output = BytesIO()
    book.save(output)
    result = parser.parse('empty-sheet.xlsx', without_sheet_dimensions(output.getvalue()))
    assert [(row['item_code'], row['quantity'], row['unit']) for row in result['data']['lines']] == [
        ('01-01-00001', '40', '码')]


@pytest.mark.parametrize('unsized', [False, True])
@pytest.mark.parametrize('sheet_count,last_row', [(1, 50001), (2, 25001)])
def test_xlsx_size_limit_applies_to_known_and_missing_dimensions(unsized, sheet_count, last_row):
    from openpyxl import Workbook
    parser = importlib.import_module('app.services.material_handover_parser')
    error = importlib.import_module('app.services.document_tools.document_ir').ToolError
    book = Workbook()
    for index in range(sheet_count):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.cell(last_row, 1, 'end')
    output = BytesIO()
    book.save(output)
    content = without_sheet_dimensions(output.getvalue()) if unsized else output.getvalue()
    with pytest.raises(error) as caught:
        parser.parse('too-large.xlsx', content)
    assert caught.value.code == 'TABLE_SIZE'


def test_file_and_numeric_boundaries():
    parser=importlib.import_module('app.services.material_handover_parser')
    error=importlib.import_module('app.services.document_tools.document_ir').ToolError
    for name,content in [('danger.xlsm',b'x'),('fake.xlsx',b'invalid'),('empty.pdf',b'')]:
        with pytest.raises(error): parser.validate_file(name,content)
    assert parser.numeric('1,000.125Y') == ('1000.125','码')
    assert parser.numeric('-12')[0] == '' and parser.numeric('1e10')[0] == ''


def test_native_pdf_table_extracts_without_fixed_template():
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    from reportlab.lib import colors
    parser=importlib.import_module('app.services.material_handover_parser')
    buffer=BytesIO(); document=SimpleDocTemplate(buffer)
    table=Table([['item code','description','qty','unit'],['00012','Black fabric','17.5','Y']],colWidths=[90,160,80,80])
    table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),1,colors.black)]));document.build([table])
    result=parser.parse('different.pdf',buffer.getvalue())
    assert len(result['data']['lines']) == 1
    assert result['data']['lines'][0]['quantity']=='17.5' and result['data']['lines'][0]['unit']=='码'


def test_uncached_formula_stays_unknown_and_displayed_code_keeps_zeroes():
    from openpyxl import Workbook
    parser=importlib.import_module('app.services.material_handover_parser')
    book=Workbook();sheet=book.active
    sheet.append(['物料编码','物料名称','数量','单位'])
    sheet.append([12,'布料','=10+20','码']);sheet['A2'].number_format='00000'
    buffer=BytesIO();book.save(buffer)
    result=parser.parse('formulas.xlsx',buffer.getvalue())
    assert result['data']['lines'][0]['item_code']=='00012'
    assert result['data']['lines'][0]['quantity']=='' and result['warnings']


def test_optional_price_columns_preserve_zero_currency_and_price_unit():
    parser=importlib.import_module('app.services.material_handover_parser')
    result=parser.parse('价格不同布局.xlsx',workbook(headers=['物料名称','单价','数量','币种','物料编码','单位','计价单位','单价依据'],
        rows=[['赠送辅料','0','15','人民币','00012','Pcs','件','免费补货'],['另种辅料','2.30','20','港币','00013','码','米','原单报价']]))
    first,second=result['data']['lines']
    assert first['unit_price']=='0' and first['currency']=='CNY' and first['price_unit']=='件'
    assert second['unit_price']=='2.30' and second['currency']=='HKD' and second['price_unit']=='米'
    assert second['quantity']=='20' and second['unit']=='码'  # no price-unit conversion
