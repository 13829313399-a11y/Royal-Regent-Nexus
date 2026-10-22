"""Mixed notes preserve source identity, authorization, material versions and atomicity."""
from uuid import uuid4
from decimal import Decimal
import pytest
from test_molding_sample_api import make_client, login_as
from test_carton_special_requests import command, create_body, created, setup, ok, accept, BASE, SUP, FACTORY
from test_carton_supplier_portal import supplier_login, accept_all

INTERNAL = "/api/carton-unified-deliveries"
EXTERNAL = "/api/carton-supplier/unified-deliveries"


def test_unified_dispatch_requires_explicit_vendor_price_for_special_lines():
    from pydantic import ValidationError
    from app.schemas.carton_unified_deliveries import Dispatch
    payload = command(delivery_note_no="MIX-MISSING-PRICE", delivery_date="2026-10-18", special_requests=[
        {"id": "request-1", "expected_revision": 1, "version_id": "version-1",
         "lines": [{"line_id": "line-1", "quantity": "2"}]}])
    with pytest.raises(ValidationError, match="不能留空"):
        Dispatch.model_validate(payload)
    payload["special_requests"][0]["lines"][0]["unit_price"] = "0"
    assert Dispatch.model_validate(payload).special_requests[0].lines[0].unit_price == 0


def test_pre_upgrade_special_retry_preserves_identity_without_ignoring_new_choices():
    import hashlib
    from types import SimpleNamespace
    from fastapi import HTTPException
    from app.schemas.carton_special_requests import Ship
    from app.services import carton_special_requests as service

    original = command(expected_revision=2, version_id="version-before-upgrade", delivery_note_no="OLD-1",
        delivery_date="2026-10-18", lines=[{"line_id": "line-1", "quantity": "2", "unit_price": "1"}])
    current = Ship.model_validate(original)
    legacy = current.model_dump(mode="json")
    legacy["lines"][0].pop("free_confirmed")
    digest = hashlib.sha256(service.pack({"action": "SHIP", "id": "request-1", "payload": legacy,
        "extra": None}).encode()).hexdigest()
    event = SimpleNamespace(detail_json=service.pack({"fingerprint": digest, "response": {"id": "saved-response"}}))
    db = SimpleNamespace(scalar=lambda statement: event)
    user = SimpleNamespace(id="supplier-1")
    assert service.replay(db, user, current, "SHIP", "request-1")[1] == {"id": "saved-response"}
    original["lines"][0]["free_confirmed"] = True
    with pytest.raises(HTTPException) as error:
        service.replay(db, user, Ship.model_validate(original), "SHIP", "request-1")
    assert error.value.status_code == 409


@pytest.fixture(autouse=True)
def isolated(monkeypatch, tmp_path):
    import test_molding_sample_api
    monkeypatch.setattr(test_molding_sample_api, "TEST_TMP_DIR", tmp_path)


def dispatch_body(row, formal=None):
    return command(delivery_note_no="MIX-" + uuid4().hex[:12], delivery_date="2026-10-18",
        formal_lines=[{"order_line_id": line["id"], "issue_id": formal["issue_id"], "quantity": "2", "unit_price": "2"}
                      for line in formal["lines"]] if formal else [],
        special_requests=[{"id": row["id"], "expected_revision": row["revision"], "version_id": row["current_version_id"],
            "lines": [{"line_id": row["lines"][0]["id"], "quantity": "10", "unit_price": "0", "free_confirmed": True}]}])


def receive_body(group, place, quantity="4"):
    detail = group["special"][0]
    shipment = detail["shipments"][0]
    return command(acceptance_date="2026-10-18", formal={"expected_revision": group["formal"]["revision"],
        "remaining_action": "WAIT", "lines": [{"shipment_line_id": line["id"], "received_quantity": "2", "unit_price": "2",
            "location_allocations": [{"location_id": place["id"], "quantity": "2"}]} for line in group["formal"]["lines"]]} if group["formal"] else None,
        special=[{"request_id": detail["id"], "shipment_id": shipment["id"], "expected_revision": detail["revision"],
            "lines": [{"shipment_line_id": shipment["lines"][0]["id"], "received_quantity": quantity,
                "location_allocations": [{"location_id": place["id"], "quantity": quantity}]}]}])


def test_mixed_dispatch_and_partial_receipt_are_one_atomic_operation(monkeypatch):
    with make_client(monkeypatch) as client:
        place = setup(client)
        row = created(client)
        supplier_login(client)
        formal = accept_all(client)
        payload = dispatch_body(row, formal)
        group = ok(client.post(EXTERNAL, json=payload), 201)
        assert group["formal"] and len(group["special"]) == 1
        assert {line["source_kind"] for line in group["lines"]} == {"FORMAL", "SPECIAL"}
        formal_lines = [line for line in group["lines"] if line["source_kind"] == "FORMAL"]
        assert all(Decimal(line["delivery_unit_price"]) == 2 for line in formal_lines)
        assert all("unit_price" not in line and line["free_confirmed"] is False for line in formal_lines)
        assert all(Decimal(line["delivery_unit_price"]) == 2 for line in group["formal"]["lines"])
        assert group["formal"]["delivery_note_no"] == group["special"][0]["shipments"][0]["delivery_note_no"]
        documents = ok(client.get("/api/carton-supplier/documents", params={"factory_id": FACTORY}))
        delivery = next(item for item in documents if item["kind"] == "DELIVERY")
        assert delivery["id"] == group["id"] and len(delivery["lines"]) == len(group["lines"])
        assert {line["order_kind"] for line in delivery["lines"]} == {"REGULAR", "SAMPLE"}
        assert {order["order_kind"] for order in delivery["orders"]} == {"REGULAR", "SAMPLE"}
        purchase = next(item for item in documents if item["kind"] == "PURCHASE")
        assert all(line["order_kind"] == "REGULAR" for line in purchase["lines"])
        from io import BytesIO
        from openpyxl import load_workbook
        exported = client.post("/api/carton-supplier/documents/export.xlsx", json={"documents": [
            {"factory_id": FACTORY, "kind": "DELIVERY", "id": group["id"]}]})
        assert exported.status_code == 200
        book = load_workbook(BytesIO(exported.content)); assert book["送货单"].max_row == len(group["lines"]) + 1; book.close()
        assert ok(client.post(EXTERNAL, json=payload), 201)["id"] == group["id"]
        collision = dict(payload, delivery_note_no="OTHER-NOTE")
        assert client.post(EXTERNAL, json=collision).status_code == 409
        login_as(client, "admin")
        pending = ok(client.get("/api/carton-supplier/internal/shipments/pending", params={"factory_id": FACTORY}))
        assert pending["total"] == 1 and pending["items"][0]["id"] == group["id"]
        body = receive_body(group, place)
        result = ok(client.post(INTERNAL + f"/{group['id']}/receive", json=body))
        assert result["status"] == "SENT"
        assert result["formal"]["status"] == "RECEIVED"
        assert result["special"][0]["shipments"][0]["lines"][0]["remaining_quantity"] == "6.0000"
        assert ok(client.post(INTERNAL + f"/{group['id']}/receive", json=body)) == result
        body2 = receive_body(result, place, "6")
        body2["formal"] = None
        done = ok(client.post(INTERNAL + f"/{group['id']}/receive", json=body2))
        assert done["status"] == "RECEIVED"
        assert ok(client.get("/api/carton-supplier/internal/shipments/pending", params={"factory_id": FACTORY}))["total"] == 0
        from sqlalchemy import select, func
        from app.db import SessionLocal
        from app.models.carton_special_requests import SpecialReceiptLink
        with SessionLocal() as db:
            assert db.scalar(select(func.count()).select_from(SpecialReceiptLink)) == 2


def test_bad_special_line_rolls_back_formal_dispatch_and_receipt(monkeypatch):
    with make_client(monkeypatch) as client:
        place = setup(client); row = created(client)
        supplier_login(client); formal = accept_all(client)
        payload = dispatch_body(row, formal)
        payload["special_requests"][0]["lines"][0]["quantity"] = "101"
        assert client.post(EXTERNAL, json=payload).status_code == 409
        assert ok(client.get(EXTERNAL, params={"factory_id": FACTORY}))["items"] == []
        payload["special_requests"][0]["lines"][0]["quantity"] = "10"
        group = ok(client.post(EXTERNAL, json=payload), 201)
        login_as(client, "admin")
        bad = receive_body(group, place, "11")
        assert client.post(INTERNAL + f"/{group['id']}/receive", json=bad).status_code == 409
        unchanged = ok(client.get(INTERNAL + f"/{group['id']}", params={"factory_id": FACTORY}))
        assert unchanged["formal"]["status"] == "SENT"
        assert not unchanged["receipts"]
        assert all(Decimal(line["received_quantity"]) == 0 for line in unchanged["lines"])


def test_marks_pending_never_gate_accepted_dispatch_and_file_evidence_is_frozen(monkeypatch):
    from PIL import Image
    from io import BytesIO
    content = BytesIO(); Image.new("RGB", (20, 20), "white").save(content, format="PNG")
    with make_client(monkeypatch) as client:
        setup(client)
        payload = create_body(); payload["lines"][0]["mark_mode"] = "PENDING"
        row = ok(client.post(BASE, json=payload), 201)
        row = ok(client.post(BASE + f"/{row['id']}/send", json=command(row)))
        supplier_login(client)
        assert client.post(EXTERNAL, json=dispatch_body(row)).status_code == 409
        draft = ok(client.post(SUP + f"/{row['id']}/attachments", data={**command(row), "role": "DRAFT", "line_id": row["lines"][0]["id"]},
            files={"file": ("mark.png", content.getvalue(), "image/png")}), 201)
        file = draft["attachments"][-1]
        supplier_try = client.post(BASE + f"/{row['id']}/lines/{row['lines'][0]['id']}/mark", json=command(draft,
            mark_mode="UPLOAD", attachment_id=file["id"]))
        assert supplier_try.status_code == 403
        login_as(client, "admin")
        row = ok(client.post(BASE + f"/{row['id']}/lines/{row['lines'][0]['id']}/mark", json=command(draft,
            mark_mode="UPLOAD", attachment_id=file["id"])))
        assert not row["lines"][0]["mark_ready"]
        assert row['current_version_id'] == draft['current_version_id']
        row = accept(client, row)
        supplier_login(client)
        group = ok(client.post(EXTERNAL, json=dispatch_body(row)), 201)
        login_as(client, "admin")
        latest = group["special"][0]
        changed = ok(client.post(BASE + f"/{row['id']}/lines/{row['lines'][0]['id']}/mark", json=command(latest, mark_mode="PENDING")))
        assert not changed["lines"][0]["mark_ready"]
        frozen = ok(client.get(INTERNAL + f"/{group['id']}", params={"factory_id": FACTORY}))["lines"][0]
        assert frozen['file_manifest'][0]['asset_id'] == draft['mark_assets'][0]['id']
        assert frozen['file_manifest'][0]['sha256'] == file['sha256']
        assert frozen['material_manifest'] == []
        original = client.get(BASE + f"/{row['id']}/attachments/{frozen['file_manifest'][0]['attachment_id']}", params={"factory_id": FACTORY})
        assert original.status_code == 200 and original.content == content.getvalue()


def test_existing_library_selection_checks_scope_revision_and_freezes_original(monkeypatch):
    with make_client(monkeypatch) as client:
        setup(client); row = created(client)
        from app.db import SessionLocal
        from app.models.carton_mark import CartonMarkAsset
        with SessionLocal() as db:
            for factory, identifier in [(FACTORY, "asset-local"), ("huakang-a", "asset-other")]:
                db.add(CartonMarkAsset(id=identifier, factory_id=factory, file_name="mark.pdf", kind="pdf", content_type="application/pdf",
                    size_bytes=8, sha256="a" * 64, content=b"original", contract_number="", recognition_source="", candidates_json="[]",
                    warning="", revision=2, is_archived=False, created_by="user-admin", created_by_name="admin", created_at="2026-10-19", updated_at="2026-10-19"))
            db.commit()
        path = BASE + f"/{row['id']}/lines/{row['lines'][0]['id']}/mark"
        assert client.post(path, json=command(row, mark_mode="EXISTING", asset_id="asset-other", asset_revision=2)).status_code == 404
        assert client.post(path, json=command(row, mark_mode="EXISTING", asset_id="asset-local", asset_revision=1)).status_code == 409
        bound = ok(client.post(path, json=command(row, mark_mode="EXISTING", asset_id="asset-local", asset_revision=2)))
        # Association advances only the library metadata and retains DRAFT bytes.
        assert bound['mark_assets'][0]['revision'] == 3
        assert bound['current_version_id'] == row['current_version_id']
        file_id = bound['attachments'][-1]['id']
        with SessionLocal() as db:
            asset = db.get(CartonMarkAsset, "asset-local"); asset.is_archived = True; asset.content = b"updated"; db.commit()
        supplier_login(client)
        assert client.get(BASE + "/mark-options", params={"factory_id": FACTORY}).status_code == 403
        assert client.get(SUP + f"/{row['id']}/attachments/{file_id}", params={"factory_id": FACTORY}).content == b"original"


def test_same_note_accepts_multiple_special_requests_without_duplicate_creation(monkeypatch):
    with make_client(monkeypatch) as client:
        setup(client); first = created(client); second = created(client)
        supplier_login(client)
        payload = dispatch_body(first)
        payload["special_requests"].extend(dispatch_body(second)["special_requests"])
        group = ok(client.post(EXTERNAL, json=payload), 201)
        assert len(group["special"]) == 2 and len(group["lines"]) == 2
        assert len({entry["shipments"][0]["delivery_note_no"] for entry in group["special"]}) == 1
        assert ok(client.get(EXTERNAL, params={"factory_id": FACTORY}))["total"] == 1
        duplicate = dispatch_body(first)
        duplicate["delivery_note_no"] = payload["delivery_note_no"]
        assert client.post(EXTERNAL, json=duplicate).status_code == 409


def test_formal_only_group_honors_current_factory_explicit_receipt_deny(monkeypatch):
    with make_client(monkeypatch) as client:
        place = setup(client); supplier_login(client); formal = accept_all(client)
        group = ok(client.post(EXTERNAL, json=command(delivery_note_no="DENY-NOTE", delivery_date="2026-10-18",
            formal_lines=[{"order_line_id": line["id"], "issue_id": formal["issue_id"], "quantity": "2", "unit_price": "2"}
                          for line in formal["lines"]])), 201)
        from test_molding_sample_api import ensure_test_user
        from test_carton_supplier_portal import grant_supplier_permissions
        ensure_test_user("engineer")
        grant_supplier_permissions("user-engineer", permissions=("carton_procurement:read", "carton_procurement:receipt_write", "carton_procurement:inventory_write"))
        login_as(client, "engineer")
        assert client.get(INTERNAL, params={"factory_id": FACTORY}).status_code == 200
        from app.db import SessionLocal
        from app.models.auth import AuthPermission, AuthUserPermissionOverride
        from sqlalchemy import select
        with SessionLocal() as db:
            permission = db.scalar(select(AuthPermission).where(AuthPermission.code == "carton_procurement:receipt_write"))
            db.add(AuthUserPermissionOverride(id="unified-receipt-deny", user_id="user-engineer", permission_id=permission.id,
                effect="deny", factory_id=FACTORY, department="*", status="active", reason="测试本厂显式拒绝"))
            db.commit()
        login_as(client, "engineer")
        assert client.get(INTERNAL, params={"factory_id": FACTORY}).status_code == 403
        body = command(acceptance_date="2026-10-18", formal={"expected_revision": group["formal"]["revision"],
            "lines": [{"shipment_line_id": line["id"], "received_quantity": "2", "unit_price": "2",
                "location_allocations": [{"location_id": place["id"], "quantity": "2"}]} for line in group["formal"]["lines"]]})
        assert client.post(INTERNAL + f"/{group['id']}/receive", json=body).status_code == 403


@pytest.mark.parametrize("received_kind", ["SAMPLE", "WORKSHOP_MISC"])
def test_extra_goods_require_warehouse_classification_and_price_zero_confirmation(monkeypatch, received_kind):
    with make_client(monkeypatch) as client:
        place = setup(client); supplier_login(client)
        payload = command(delivery_note_no="EXTRA-1", delivery_date="2026-10-18", extra_lines=[dict(
            create_body()["lines"][0], unit_price="0", purpose="周转使用", free_confirmed=False)])
        assert client.post(EXTERNAL, json=payload).status_code == 422
        payload["extra_lines"][0]["free_confirmed"] = True
        group = ok(client.post(EXTERNAL, json=payload), 201)
        assert group["special"][0]["origin"] == "SUPPLIER_EXTRA"
        documents = ok(client.get("/api/carton-supplier/documents", params={"factory_id": FACTORY}))
        extra_document = next(item for item in documents if item["document_no"] == "EXTRA-1")
        assert {line["order_kind"] for line in extra_document["lines"]} == {"UNKNOWN"}
        login_as(client, "admin")
        body = receive_body(group, place)
        assert client.post(INTERNAL + f"/{group['id']}/receive", json=body).status_code == 422
        body["special"][0].update(received_kind=received_kind)
        received = ok(client.post(INTERNAL + f"/{group['id']}/receive", json=body))
        assert received["special"][0]["kind"] == received_kind
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonReceiptLine
        from sqlalchemy import select
        with SessionLocal() as db:
            assert db.scalar(select(CartonReceiptLine.customer_code)) == f"__{received_kind}__"
        supplier_login(client)
        documents = ok(client.get("/api/carton-supplier/documents", params={"factory_id": FACTORY}))
        extra_document = next(item for item in documents if item["document_no"] == "EXTRA-1")
        assert {line["order_kind"] for line in extra_document["lines"]} == {"SAMPLE" if received_kind == "SAMPLE" else "MISC"}
