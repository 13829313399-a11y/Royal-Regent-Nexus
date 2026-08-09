from __future__ import annotations

from datetime import datetime
from io import BytesIO

from app.db import Base
from app.models import (
    injection_scheduling as _injection_scheduling_models,  # noqa: F401
)
from app.models.injection_scheduling_execution import InjectionSchedulingAuditEvent
from app.models.injection_scheduling_import import (
    InjectionSchedulingImportProfile,
    InjectionSchedulingImportProfileFactory,
)
from app.services.auth import AuthContext
from app.services.injection_scheduling_excel import parse_injection_scheduling_workbook
from app.services.injection_scheduling_profile_registry import (
    active_profiles_for_factory,
    create_profile_revision,
    seed_builtin_import_profiles,
    transition_profile,
)
from app.services.injection_scheduling_profiles import (
    BUILTIN_IMPORT_PROFILES,
    CANONICAL_FIELD_CATALOG,
    normalize_header,
    profile_config,
    profile_from_config,
)
from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session


def _bytes(workbook: Workbook) -> bytes:
    payload = BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def _huakang_b_fixture(*, drift_order_header: bool = False) -> bytes:
    workbook = Workbook()
    plan = workbook.active
    plan.title = "排期表"
    plan["A1"] = "华康B啤机生产日排表（脱敏契约样本）"
    for column, value in {
        "A": "机位",
        "B": "机位",
        "E": "货号",
        "G": "工模",
        "H": "名称",
        "I": "漂移单据" if drift_order_header else "单号",
        "J": "仓库",
        "L": "订单数",
        "M": "已啤数",
        "O": "计划日目标",
        "P": "用料",
        "Q": "水口比例",
        "R": "颜色",
        "S": "色粉",
        "X": "下单期",
        "Y": "开始交货期",
        "Z": "交货完成期",
        "AE": "计划啤货期",
        "AF": "计划完成期",
        "AK": "是否喷油",
        "AN": "每班计划啤数",
        "AT": "备注",
        "AU": "单双臂",
        "AV": "气剪",
        "AW": "白班 8/5",
        "AX": "夜班 8/5",
    }.items():
        plan[f"{column}3"] = value

    plan["A4"] = "B01"
    plan["B4"] = "B01"

    for column, value in {
        "B": "B01",
        "E": "ITEM-01",
        "G": "MOLD-01",
        "H": "脱敏产品甲",
        "I": "ORDER-01",
        "J": "WH-B",
        "L": 100,
        "M": 10,
        "O": 200,
        "P": "PP",
        "Q": "5%",
        "R": "黑",
        "S": "C-01",
        "AE": datetime(2026, 8, 5, 8),  # noqa: DTZ001 - Excel stores local wall time
        "AF": datetime(2026, 8, 5, 20),  # noqa: DTZ001 - Excel stores local wall time
        "AN": 50,
        "AW": 50,
    }.items():
        plan[f"{column}5"] = value

    # No assignment/window means backlog even though it follows a machine block.
    for column, value in {
        "E": "ITEM-02",
        "G": "MOLD-02",
        "H": "脱敏产品乙",
        "I": "ORDER-02",
        "L": 50,
        "M": 60,
        "AN": 25,
    }.items():
        plan[f"{column}6"] = value

    machine = workbook.create_sheet("厂区现有啤机")
    machine["A3"] = "设备编号"
    machine["B3"] = "设备名称"
    machine["A4"] = "B01"
    machine["B4"] = "320T 注塑机"
    machine["D4"] = "32A"
    machine["F4"] = 500
    machine["G4"] = "700x700"
    machine["H4"] = 320
    machine["I4"] = "标准机"
    machine["J4"] = "双臂"

    for title in ("已啤完", "机台完成时间", "外发2", "取消订单", "每班啤数记录"):
        history = workbook.create_sheet(title)
        history["I5"] = "HISTORY-MUST-NOT-BECOME-CURRENT"
        history["L5"] = 999
    return _bytes(workbook)


def _huaxing_fixture() -> bytes:
    workbook = Workbook()
    plan = workbook.active
    plan.title = "计划表"
    plan["A1"] = "河源华兴啤机生产日计划表（脱敏契约样本）"
    for column, value in {
        "B": "机号",
        "G": "工模",
        "H": "名称",
        "I": "单号",
        "J": "货号",
        "L": "订单数",
        "M": "已啤数",
        "O": "计划目标",
        "AG": "计划生产期",
        "AH": "计划完成期",
        "AR": "仓库",
    }.items():
        plan[f"{column}3"] = value
    for column, value in {
        "B": "HX01",
        "G": "MOLD-HX-01",
        "H": "脱敏产品",
        "I": "HX-ORDER-01",
        "J": "HX-ITEM-01",
        "L": 80,
        "M": 8,
        "O": 40,
        "AG": datetime(2026, 8, 5, 8),  # noqa: DTZ001 - Excel stores local wall time
        "AH": datetime(2026, 8, 5, 20),  # noqa: DTZ001 - Excel stores local wall time
    }.items():
        plan[f"{column}5"] = value
    return _bytes(workbook)


def test_header_normalization_and_controlled_catalog_contract():
    assert normalize_header(" 订单数（PCS）\n") == "订单数"
    assert normalize_header("每班-计划：啤数") == "每班计划啤数"
    assert CANONICAL_FIELD_CATALOG["planned_start"].authority == "BASELINE_DECISION"
    assert (
        CANONICAL_FIELD_CATALOG["source_outstanding_quantity"].authority
        == "SYSTEM_DERIVED"
    )


def test_huakang_b_profile_maps_duplicate_header_by_coordinate_and_classifies_rows():
    normalized, issues = parse_injection_scheduling_workbook(
        _huakang_b_fixture(),
        "huakang-b-contract.xlsx",
        factory_id="huakang-b",
        system_machine_codes={"B01"},
        system_mold_nos={"MOLD-01", "MOLD-02"},
    )

    assert normalized["profile"]["profile_code"] == "huakang_b_daily_plan_v1"
    assert normalized["batch_state"] == "PREVIEW_READY"
    assert normalized["summary"]["scheduled_baseline_count"] == 1
    assert normalized["summary"]["backlog_count"] == 1
    assert normalized["summary"]["order_count"] == 2
    assert normalized["summary"]["master_difference_count"] == 0
    assert normalized["scheduled_baseline_tasks"][0]["machine_code"] == "B01"
    assert normalized["scheduled_baseline_tasks"][0]["daily_target_quantity"] == 200
    assert normalized["scheduled_baseline_tasks"][0]["shift_target_quantity"] == 50
    machine_mapping = next(
        item
        for item in normalized["mapping"]
        if item["canonical_field"] == "machine_code"
    )
    assert machine_mapping["column"] == "B"
    assert machine_mapping["raw_header"] == "机位"
    assert normalized["optional_completed_history"][0]["persisted"] is False
    assert all(
        order["order_no"] != "HISTORY-MUST-NOT-BECOME-CURRENT"
        for order in normalized["orders"]
    )
    overproduction = [item for item in issues if item["code"] == "OVERPRODUCED"]
    assert len(overproduction) == 1
    assert overproduction[0]["blocking"] is False


def test_huaxing_and_unknown_factory_use_same_canonical_contract():
    source = _huaxing_fixture()
    normalized, _ = parse_injection_scheduling_workbook(
        source,
        "huaxing-contract.xlsx",
        factory_id="huaxing",
        system_machine_codes={"HX01"},
        system_mold_nos={"MOLD-HX-01"},
    )
    assert normalized["profile"]["profile_code"] == "huaxing_daily_plan_v1"
    assert normalized["summary"]["scheduled_baseline_count"] == 1
    assert normalized["schema_version"] == "injection-scheduling-canonical-v1"

    unknown, issues = parse_injection_scheduling_workbook(
        source,
        "unknown-factory.xlsx",
        factory_id="huakang-c",
    )
    assert unknown["batch_state"] == "MAPPING_REQUIRED"
    assert unknown["profile"] is None
    assert {item["code"] for item in issues} == {"PROFILE_NOT_IDENTIFIED"}


def test_required_header_drift_is_mapping_required_and_not_confirmable():
    normalized, issues = parse_injection_scheduling_workbook(
        _huakang_b_fixture(drift_order_header=True),
        "huakang-b-drift.xlsx",
        factory_id="huakang-b",
        system_mold_nos={"MOLD-01", "MOLD-02"},
    )
    assert normalized["batch_state"] == "MAPPING_REQUIRED"
    assert normalized["summary"]["can_confirm"] is False
    order_mapping = next(
        item for item in normalized["mapping"] if item["canonical_field"] == "order_no"
    )
    assert order_mapping["status"] == "MISSING"
    assert any(item["code"] == "REQUIRED_MAPPING_MISSING" for item in issues)


def test_builtin_profile_registry_is_idempotent_and_factory_scoped():
    engine = create_engine("sqlite://", future=True)
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        seed_builtin_import_profiles(db)
        seed_builtin_import_profiles(db)
        assert (
            db.scalar(
                select(func.count()).select_from(InjectionSchedulingImportProfile)
            )
            == 4
        )
        assert (
            db.scalar(
                select(func.count()).select_from(
                    InjectionSchedulingImportProfileFactory
                )
            )
            == 14
        )
        assert (
            db.scalar(select(func.count()).select_from(InjectionSchedulingAuditEvent))
            == 14
        )
        huakang_profiles = active_profiles_for_factory(db, "huakang-b")
        assert [item.profile_code for item in huakang_profiles] == [
            "huakang_b_daily_plan_v1"
        ]
        assert active_profiles_for_factory(db, "huakang-c") == ()
        assert [
            item.profile_code
            for item in active_profiles_for_factory(
                db, "huakang-c", document_kind="DEMAND_ORDER"
            )
        ] == ["demand_order_shared_v1"]
        record = db.get(
            InjectionSchedulingImportProfile, "isprofile-huakang-b-daily-v1"
        )
        config = profile_from_config(__import__("json").loads(record.config_json))
        assert config.revision == record.revision == 1
        assert config.status == "ACTIVE"


def test_profile_revision_activation_retires_predecessor_and_audits():
    engine = create_engine("sqlite://", future=True)
    Base.metadata.create_all(engine)
    actor = AuthContext(
        id="user-profile-reviewer",
        username="profile-reviewer",
        display_name="模板审核员",
        roles=("molding_supervisor",),
        role_codes=("molding_supervisor",),
        permissions=frozenset({"injection_scheduling:manage_import_profiles"}),
        factory_scopes=("huakang-b",),
        department_scopes=("molding",),
    )
    reviewer = AuthContext(
        id="user-profile-approver",
        username="profile-approver",
        display_name="另一位模板审核员",
        roles=("molding_supervisor",),
        role_codes=("molding_supervisor",),
        permissions=frozenset({"injection_scheduling:manage_import_profiles"}),
        factory_scopes=("huakang-b",),
        department_scopes=("molding",),
    )
    with Session(engine) as db:
        seed_builtin_import_profiles(db)
        source = next(
            item
            for item in BUILTIN_IMPORT_PROFILES
            if item.profile_family == "huakang_b_daily_plan"
        )
        draft = create_profile_revision(
            db,
            factory_id="huakang-b",
            profile_family="huakang_b_daily_plan",
            profile_code="huakang_b_daily_plan_v2",
            name="华康 B 啤机日排表 v2",
            description="脱敏契约测试修订",
            expected_family_revision=1,
            request_id="profile-create-v2",
            config=profile_config(source),
            user=reviewer,
        )
        assert draft.revision == 2
        assert draft.status == "PROFILE_DRAFT"
        activated = transition_profile(
            db,
            profile_id=draft.id,
            factory_id="huakang-b",
            expected_lifecycle_revision=1,
            target_status="ACTIVE",
            request_id="profile-activate-v2",
            reason="真实样本映射契约复核通过",
            user=actor,
        )
        assert activated.status == "ACTIVE"
        predecessor = db.get(
            InjectionSchedulingImportProfile, "isprofile-huakang-b-daily-v1"
        )
        assert predecessor.status == "RETIRED"
        assert [
            item.revision for item in active_profiles_for_factory(db, "huakang-b")
        ] == [2]
        event_types = set(
            db.scalars(select(InjectionSchedulingAuditEvent.event_type)).all()
        )
        assert {
            "import_profile_revision_created",
            "import_profile_retired_by_successor",
            "import_profile_activated",
        } <= event_types
