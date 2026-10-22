"""Confirmed month evidence freezes; reopening does not repost stock or money."""
import json
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from fastapi import HTTPException
from sqlalchemy import func, select

from test_molding_sample_api import make_client, login_as
from test_carton_supplier_portal import (
    BASE as PORTAL, setup_portal, supplier_login, accept_all, ship_payload, receive_payload,
    _dongkang_delivery_file,
)
from test_carton_settlement_collaboration import workspace, save, review
from test_carton_supplier_settlement_api import action
from test_carton_special_requests import BASE as SPECIAL, SUP, command, created
from test_carton_procurement_api import _order_payload, _submit_order


@pytest.fixture(autouse=True)
def isolated_database(monkeypatch, tmp_path):
    import test_molding_sample_api
    monkeypatch.setattr(test_molding_sample_api, "TEST_TMP_DIR", tmp_path)


def test_evidence_guard_validates_all_dates_and_deduplicates_periods(monkeypatch):
    from app.services import carton_supplier_settlement as settlement
    from app.services.carton_supplier_evidence_guard import ensure_evidence_open
    checked = []
    monkeypatch.setattr(settlement, "ensure_open", lambda db, factory, supplier, period:
        checked.append((factory, supplier, period)))
    ensure_evidence_open(None, "factory", "supplier", "2026-09-21", date(2026, 9, 30), "2026-10")
    assert checked == [("factory", "supplier", "2026-09"), ("factory", "supplier", "2026-10")]
    for invalid in (None, "", "2026-13", "2026-02-30", "20260921", "2026-9-21", "2026-09-anything"):
        with pytest.raises(HTTPException) as error:
            ensure_evidence_open(None, "factory", "supplier", "2026-09-21", invalid)
        assert error.value.status_code == 422
    with pytest.raises(HTTPException) as error:
        ensure_evidence_open(None, "factory", "supplier")
    assert error.value.status_code == 422


def quantities(db):
    from app.models.carton_supplier_portal import SupplierShipment
    from app.models.carton_special_requests import SpecialShipment, SpecialRequest, UnifiedDelivery
    from app.models.carton_procurement import CartonReceipt, CartonInventoryMovement
    return tuple(db.scalar(select(func.count()).select_from(model)) for model in (
        SupplierShipment, SpecialShipment, SpecialRequest, UnifiedDelivery, CartonReceipt, CartonInventoryMovement))


def source_counts():
    from app.db import SessionLocal
    with SessionLocal() as db:
        return quantities(db)


def blocked(response):
    assert response.status_code == 409, response.text
    assert "2026-09" in response.text and "重开" in response.text, response.text


def delivery_file(order, notes):
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(_dongkang_delivery_file(order, source_price=2)))
    sheet = workbook.active
    original = list(sheet.values)[1:]
    sheet.delete_rows(2, sheet.max_row)
    for note, day, quantity in notes:
        for values in original:
            values = list(values)
            values[0], values[1], values[7] = note, day, quantity
            sheet.append(values)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return {"file": ("原送货.xlsx", output.getvalue())}


def test_confirmed_evidence_rejects_backfill_but_allows_next_month_and_reasoned_reopen(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client)
        from app.services import carton_procurement as core
        monkeypatch.setattr(core, "business_now", lambda: datetime(2026, 11, 5, 12))
        order = accept_all(client)
        special = created(client)
        supplier_login(client)
        original = ship_payload(order)
        for paper in original["lines"]:
            paper["unit_price"] = "2"
        response = client.post(PORTAL + "/shipments", json=original)
        assert response.status_code == 201, response.text
        shipment = response.json()
        login_as(client, "admin")
        acceptance = receive_payload(client, shipment) | {"acceptance_date": "2026-09-24", "remaining_action": "WAIT"}
        for line in acceptance["lines"]:
            line.update(received_quantity="2", location_allocations=[{**line["location_allocations"][0], "quantity": "2"}])
        response = client.post(PORTAL + f"/internal/shipments/{shipment['id']}/receive", json=acceptance)
        assert response.status_code == 200, response.text
        shipment = response.json()

        # Preserve an unmatched original before a genuine formal source exists.
        # Linking it later changes its frozen vendor identity, even without stock.
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(["送货单号", "日期", "客户", "客户单号", "客户料号", "名称", "纸质", "规格", "数量", "单价"])
        workbook.active.append(["UNMATCHED-SEP", "2026-09-25", "华兴", "LATE-ASSOC", "LATE-BOX", "外箱", "A33+B", "20*10*8 cm", 1, 2])
        stream = BytesIO(); workbook.save(stream); workbook.close()
        unmatched_upload = {"file": ("原无单送货.xlsx", stream.getvalue())}
        supplier_login(client)
        preview = client.post(PORTAL + "/shipments/import-preview", files=unmatched_upload).json()
        response = client.post(PORTAL + "/shipments/import-confirm", files=unmatched_upload, data={"sha256": preview["sha256"],
            "selections": json.dumps([{"factory_id": "huaxing", "delivery_note_no": "UNMATCHED-SEP"}])})
        assert response.status_code == 200, response.text
        unmatched = response.json()["shipments"][0]
        login_as(client, "admin")
        target_body = _order_payload()
        target_body.update(contract_no="LATE-ASSOC", item_no="LATE-BOX")
        target_body["lines"] = [{**target_body["lines"][0], "specification": "20*10*8 cm", "dimension_unit": "cm"}]
        response = client.post("/api/carton-procurement/orders", json=target_body)
        assert response.status_code == 201, response.text
        target = _submit_order(client, response.json())
        pending_link = {"factory_id": "huaxing", "expected_revision": unmatched["revision"],
            "order_line_id": target["lines"][0]["id"], "expected_order_revision": target["revision"],
            "customer_code": "DICKIE", "reason": "原单纸品后续补齐正式来源"}
        pending_link_url = PORTAL + f"/internal/shipments/{unmatched['id']}/lines/{unmatched['lines'][0]['id']}/link-order"
        september = workspace(client)
        saved = save(client, september)
        supplier_login(client)
        approved = review(client, saved)
        assert approved.status_code == 200, approved.text
        login_as(client, "admin")
        response = action(client, approved.json())
        assert response.status_code == 200, response.text
        confirmed = response.json()
        assert confirmed["status"] == "CONFIRMED"
        frozen = {key: confirmed[key] for key in ("sources", "statement", "result", "source_fingerprint")}
        fingerprint = workspace(client)["collaboration"]["source_fingerprint"]

        # No new September evidence: an existing original can be accepted in
        # October, without changing the September as-of evidence or amount.
        acceptance = receive_payload(client, shipment) | {"acceptance_date": "2026-10-02", "remaining_action": "WAIT"}
        response = client.post(PORTAL + f"/internal/shipments/{shipment['id']}/receive", json=acceptance)
        assert response.status_code == 200, response.text
        assert workspace(client)["collaboration"]["source_fingerprint"] == fingerprint
        assert all(Decimal(row["quantity"]) == 8 for row in workspace(client, "2026-10")["sources"])

        # A normal later-month warehouse receipt remains available. Its old
        # original note is evidence only when the supplier subsequently uploads it.
        locations = client.get("/api/carton-procurement/inventory/locations", params={"factory_id": "huaxing"}).json()
        location = next(row["id"] for row in locations if row["bin_code"] == "A-01")
        manual = {"factory_id": "huaxing", "request_id": uuid4().hex, "post_immediately": True,
            "delivery_note_no": "LATE-ORIGINAL", "delivery_date": "2026-09-25", "acceptance_date": "2026-10-03",
            "lines": [{"order_line_id": line["id"], "delivered_quantity": "2", "received_quantity": "2", "unit_price": "2",
                "location_allocations": [{"location_id": location, "quantity": "2"}]} for line in order["lines"]]}
        response = client.post("/api/carton-procurement/receipts", json=manual)
        assert response.status_code == 201, response.text
        receipt = response.json()
        current_order = next(row for row in client.get("/api/carton-procurement/orders",
            params={"factory_id": "huaxing"}).json()["items"] if row["order_no"] == order["order_no"])
        before = source_counts()

        supplier_login(client)
        blocked(client.post(PORTAL + "/shipments", json=ship_payload(order) | {"delivery_note_no": "NEW-OLD"}))
        special_payload = command(special, version_id=special["current_version_id"], delivery_note_no="SPECIAL-OLD",
            delivery_date="2026-09-25", lines=[{"line_id": special["lines"][0]["id"], "quantity": "1", "unit_price": "2"}])
        blocked(client.post(SUP + f"/{special['id']}/ship", json=special_payload))
        extra = {"packaging_type": "外箱", "paper_quality": "A33+B", "specification": "20*10*8", "dimension_unit": "cm",
            "quantity": "1", "unit": "个", "mark_mode": "NONE", "unit_price": "2", "purpose": "实物额外到货"}
        blocked(client.post(PORTAL + "/unified-deliveries", json=command(
            delivery_note_no="EXTRA-OLD", delivery_date="2026-09-25", extra_lines=[extra])))
        assert source_counts() == before  # Also rolls back supplier-extra requests.

        batch_upload = delivery_file(order, [("A-OPEN-OCT", "2026-10-04", 1), ("Z-CLOSED-SEP", "2026-09-25", 1)])
        preview = client.post(PORTAL + "/shipments/import-preview", files=batch_upload)
        assert preview.status_code == 200, preview.text
        groups = {row["delivery_note_no"]: row for row in preview.json()["groups"]}
        assert groups["A-OPEN-OCT"]["ready"]
        assert not groups["Z-CLOSED-SEP"]["ready"]
        blocked(client.post(PORTAL + "/shipments/import-confirm", files=batch_upload, data={"sha256": preview.json()["sha256"],
            "selections": json.dumps([{"factory_id": "huaxing", "delivery_note_no": note} for note in groups])}))
        assert source_counts() == before  # October child must roll back too.
        late_upload = delivery_file(order, [("LATE-ORIGINAL", "2026-09-25", 2)])
        preview = client.post(PORTAL + "/shipments/import-preview", files=late_upload).json()
        assert not preview["groups"][0]["receipt_link_ready"]
        late_data = {"sha256": preview["sha256"], "selections": json.dumps([
            {"factory_id": "huaxing", "delivery_note_no": "LATE-ORIGINAL", "registration_mode": "EXISTING_RECEIPT"}])}
        blocked(client.post(PORTAL + "/shipments/import-confirm", files=late_upload, data=late_data))

        login_as(client, "admin")
        manual_special = command(special, delivery_note_no="MANUAL-SP-OLD", delivery_date="2026-09-25", acceptance_date="2026-10-04",
            lines=[{"line_id": special["lines"][0]["id"], "received_quantity": "1", "unit_price": "2",
                "location_allocations": [{"location_id": location, "quantity": "1"}]}])
        blocked(client.post(SPECIAL + f"/{special['id']}/manual-receive", json=manual_special))
        manual_formal = command(delivery_note_no="MANUAL-OLD", delivery_date="2026-09-25", acceptance_date="2026-10-04",
            order_revisions={order["order_no"]: current_order["revision"]}, formal_lines=manual["lines"])
        blocked(client.post("/api/carton-unified-deliveries/warehouse-receive", json=manual_formal))
        assert source_counts() == before
        current = workspace(client)
        assert current["collaboration"]["source_fingerprint"] == fingerprint
        assert {key: current["documents"][0][key] for key in frozen} == frozen
        blocked(client.post(pending_link_url, json=pending_link))

        # A confirmed currency in the same supplier/month freezes the shared
        # evidence even if it is not the usual CNY. Factory/supplier stay scoped.
        from app.db import SessionLocal
        from app.models.carton_supplier_settlement import CartonSupplierSettlement
        from app.services.carton_supplier_evidence_guard import ensure_evidence_open
        with SessionLocal() as db:
            extra_doc = CartonSupplierSettlement(id="test-evidence-usd", factory_id="huaxing", supplier_id=current["supplier_id"],
                period="2026-08", currency="USD", version=1, revision=1, status="CONFIRMED", source_fingerprint="a" * 64,
                created_at="2026-09-01", updated_at="2026-09-01", created_by="test")
            db.add(extra_doc); db.commit()
            with pytest.raises(HTTPException) as error:
                ensure_evidence_open(db, "huaxing", current["supplier_id"], "2026-08-31", "2026-10-04")
            assert error.value.status_code == 409 and "2026-08" in error.value.detail
            ensure_evidence_open(db, "huaxing", "another-supplier", "2026-08-31")
            ensure_evidence_open(db, "another-factory", current["supplier_id"], "2026-08-31")

        reopened = action(client, confirmed, "reopen", reason="后补九月原送货凭证重新核实")
        assert reopened.status_code == 200, reopened.text
        assert reopened.json()["status"] == "DRAFT" and "supplier_review" not in reopened.json()["result"]
        response = client.post(pending_link_url, json=pending_link)
        assert response.status_code == 200, response.text
        supplier_login(client)
        response = client.post(PORTAL + "/shipments/import-confirm", files=late_upload, data=late_data)
        assert response.status_code == 200, response.text
        late = response.json()["shipments"][0]
        login_as(client, "admin")
        link = {"factory_id": "huaxing", "request_id": uuid4().hex, "expected_revision": late["revision"],
            "receipt_id": receipt["id"], "expected_receipt_revision": receipt["revision"], "reason": "后补原单核实既有入库"}
        response = client.post(PORTAL + f"/internal/shipments/{late['id']}/link-receipt", json=link)
        assert response.status_code == 200, response.text
        assert source_counts()[-2:] == before[-2:]  # No receipt or inventory repost.
        refreshed = workspace(client)
        assert refreshed["collaboration"]["source_fingerprint"] != fingerprint
        updated = save(client, refreshed, doc=reopened.json())
        assert action(client, updated).status_code == 409  # Must obtain fresh supplier confirmation.
        supplier_login(client)
        approved = review(client, updated)
        assert approved.status_code == 200, approved.text
        login_as(client, "admin")
        response = action(client, approved.json())
        assert response.status_code == 200, response.text
        final = response.json()
        assert final["result"]["net_amount"] == confirmed["result"]["net_amount"]
        history = workspace(client)["documents"]
        assert history[1]["status"] == "SUPERSEDED"
        assert {key: history[1][key] for key in frozen} == frozen
        supplier_login(client)
        # Successful requests replay after the month closes again; no new evidence.
        assert client.post(PORTAL + "/shipments", json=original).status_code == 201
        assert client.post(PORTAL + "/shipments/import-confirm", files=late_upload, data=late_data).status_code == 200
        login_as(client, "admin")
        assert client.post(PORTAL + f"/internal/shipments/{late['id']}/link-receipt", json=link).status_code == 200
        assert client.post(pending_link_url, json=pending_link).status_code == 200
