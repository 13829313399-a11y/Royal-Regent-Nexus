"""Original layout, QR readability, immutable evidence and fresh public/staff scope."""
import copy
import importlib
import io
import json
from pathlib import Path
from uuid import uuid4

import pytest
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.pagesizes import A4, landscape
from sqlalchemy import select

from test_material_handover import client, issuer, confirmed, upload, workspace, BASE, FACTORY


def data():
    from app.schemas.material_handover import HandoverData, HandoverLine
    return HandoverData(kind='DELIVERY', source_no='XS202610080001', document_date='2026-10-08', counterparty='供应商原单样张',
        lines=[HandoverLine(id='L1', item_code='01-01-10-00002', item_name='58寸黑色 210T/TT18', quantity='2000', unit='码', lot='1#',
                          roll_quantities='200 / 200 / 200 / 200 / 200 / 200 / 200 / 200 / 200 / 200')]).model_dump(mode='json')


def pdf_source(pages=1):
    from app.services.material_handover_print import font_name
    stream = io.BytesIO(); canvas = Canvas(stream, pagesize=landscape(A4))
    for index in range(pages):
        width, height = landscape(A4)
        canvas.setFont(font_name(), 19)
        canvas.drawCentredString(width/2, height-32, '供应商原来的销售出货单')
        canvas.setFont(font_name(), 11)
        canvas.drawString(24, height-61, '原单左上角文字也要保留，不能用二维码直接覆盖')
        canvas.drawRightString(width-24, height-61, '单号 XS202610080001')
        canvas.drawString(24, height-86, '日期 2026-10-08 / 客户 华康塑胶制品 / 物料 01-01-10-00002')
        canvas.rect(24, 72, width-48, height-190)
        canvas.line(24, height-147, width-24, height-147)
        canvas.drawString(38, height-130, '原单版式：订单号 / 布料名称 / 缸号 / 每卷数量 / 合计数量')
        canvas.drawString(38, height-170, 'CGDD014447 / 58寸黑色 210T/TT18 / 1# / 200 × 10卷 / 2000码')
        canvas.drawString(24, 44, f'原单签收区及备注必须保留 - 第 {index+1} 页')
        canvas.showPage()
    canvas.save(); return stream.getvalue()


@pytest.mark.parametrize('corner', ['left', 'right'])
def test_original_page_all_content_is_retained_below_the_qr_band_and_code_decodes(tmp_path, corner):
    import pdfplumber
    import pypdfium2 as pdfium
    import cv2
    import numpy as np
    from app.services.material_handover_print import render
    raw = pdf_source(2); before = raw; body = data(); code = 'MH1:MHD-'+'a'*32+':1'
    result = render('原单.pdf', raw, body, {'data': body}, code, corner, tmp_path)
    assert raw == before
    with pdfplumber.open(io.BytesIO(result)) as text_pdf, pdfium.PdfDocument(result) as output:
        assert len(output) == 2
        for index in range(len(output)):
            page = output[index]
            assert tuple(page.get_size()) == pytest.approx(tuple(map(float, landscape(A4))), abs=.005)
            images = text_pdf.pages[index].images
            assert len(images) == 1 and images[0]['top'] >= 100  # header cannot cover original paper
            bitmap = page.render(scale=3); picture=bitmap.to_pil().convert('RGB')
            image = np.asarray(picture)
            decoded, points, _ = cv2.QRCodeDetector().detectAndDecode(image[:300])
            assert decoded == code and points is not None
            center = points.mean(axis=1)[0][0]
            assert (center < picture.width/2) == (corner == 'left')
            picture.close(); bitmap.close(); page.close()
        # No action, external source URL or issuer credential enters the static PDF.
        assert 'http' not in ''.join(p.extract_text() or '' for p in text_pdf.pages)


def test_changed_confirmed_quantities_attach_the_current_values_without_rewriting_original(tmp_path):
    from pypdf import PdfReader
    from app.services.material_handover_print import render
    body = data(); edited = copy.deepcopy(body)
    edited['lines'][0].update(quantity='1800', unit_price='2.3', price_unit='码', currency='CNY', price_note='核对后的单价依据')
    result = render('原单.pdf', pdf_source(), edited, {'data': body}, 'MH1:MHD-test:1', 'right', tmp_path)
    output = PdfReader(io.BytesIO(result))
    assert len(output.pages) == 2
    text = output.pages[1].extract_text()
    assert '1800' in text and '2.3' in text and 'CNY' in text and '系统确认明细' in text
    assert '核对后的单价依据' in text


@pytest.mark.parametrize('suffix', ['xlsx', 'docx'])
def test_office_original_is_rendered_as_a_print_copy_with_original_tables(tmp_path, suffix):
    import pdfplumber
    from app.services.material_handover_print import render
    from app.services.document_tools.office_engine import office_executable, uno_python
    if not office_executable() or not uno_python(): pytest.skip('Office/UNO renderer not installed')
    if suffix == 'xlsx':
        from openpyxl import Workbook
        book = Workbook(); sheet = book.active
        sheet.append(['SUPPLIER ORIGINAL DELIVERY NOTE']); sheet.append(['物料编码','物料名称','缸号','数量','单位'])
        sheet.append(['000123','同种布料','1#',2000,'码'])
        sheet.print_area='A1:E3'; sheet.page_setup.orientation='landscape'; sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
        stream = io.BytesIO(); book.save(stream)
    else:
        from docx import Document
        doc = Document(); doc.add_heading('SUPPLIER ORIGINAL DELIVERY NOTE',0)
        table = doc.add_table(rows=2, cols=5)
        for cell,value in zip(table.rows[0].cells, ['物料编码','物料名称','缸号','数量','单位']): cell.text=value
        for cell,value in zip(table.rows[1].cells, ['000123','同种布料','1#','2000','码']): cell.text=value
        stream=io.BytesIO(); doc.save(stream)
    body=data()
    result=render('original.'+suffix, stream.getvalue(), body, {'data':body}, 'MH1:MHD-test:1', 'right', tmp_path)
    with pdfplumber.open(io.BytesIO(result)) as output:
        assert len(output.pages) >= 1 and output.pages[0].images


def test_external_office_resources_and_signed_pdf_do_not_reach_rendering(tmp_path):
    from app.services.material_handover_print import render
    from app.services.document_tools.document_ir import ToolError
    from zipfile import ZipFile
    stream=io.BytesIO()
    with ZipFile(stream,'w') as archive:
        archive.writestr('word/_rels/document.xml.rels','<Relationships><Relationship Type="image" TargetMode="External" Target="file:///private/file"/></Relationships>')
    body=data()
    with pytest.raises(ToolError, match='外部'):
        render('external.docx',stream.getvalue(),body,{'data':body},'MH1:MHD-test:1','right',tmp_path)
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import NameObject, DictionaryObject
    writer=PdfWriter();writer.append_pages_from_reader(PdfReader(io.BytesIO(pdf_source())))
    writer._root_object[NameObject('/Perms')]=DictionaryObject({NameObject('/DocMDP'):DictionaryObject()})
    stream=io.BytesIO();writer.write(stream)
    with pytest.raises(ToolError,match='签名'):
        render('signed.pdf',stream.getvalue(),body,{'data':body},'MH1:MHD-test:1','right',tmp_path)


def test_print_read_is_scoped_confirmed_and_never_posts_or_changes_original(client, monkeypatch):
    party,headers=issuer(client);_,other=issuer(client,'别的供方')
    draft=upload(client,headers)
    files=importlib.import_module('app.services.material_handover_files')
    calls=[]
    monkeypatch.setattr(files,'render_original',lambda *args: calls.append(args) or b'%PDF-1.4\nprint-copy')
    url=BASE+'/external/documents/'+draft['id']+'/print'
    assert client.get(url,headers=headers,params={'expected_revision':draft['revision']}).status_code==409
    assert not calls
    body=copy.deepcopy(draft['data']);body.update(source_no='PRINT-ORIGINAL',document_date='2026-07-18')
    response=client.post(BASE+'/external/documents/'+draft['id']+'/confirm',headers=headers,
        json=dict(request_id=str(uuid4()),expected_revision=draft['revision'],acknowledge_warnings=True,data=body))
    assert response.status_code==200,response.text
    doc=response.json();before=client.get(BASE+'/external/documents/'+doc['id']+'/source',headers=headers).content
    assert client.get(url,headers=other,params={'expected_revision':doc['revision']}).status_code==404
    assert not calls
    assert client.get(url,headers=headers,params={'expected_revision':doc['revision']-1}).status_code==409
    response=client.get(url,headers=headers,params={'expected_revision':doc['revision'],'corner':'left'})
    assert response.status_code==200 and response.headers['content-type']=='application/pdf'
    assert response.headers['cache-control']=='private, no-store'
    assert calls[0][-1]=='left' and calls[0][-2]==doc['code']
    assert client.get(BASE+'/external/documents/'+doc['id'],headers=headers).json()==doc
    assert client.get(BASE+'/external/documents/'+doc['id']+'/source',headers=headers).content==before
    assert workspace(client)['stock']==[]
    assert client.get(BASE+'/documents/'+doc['id']+'/print',params={'factory_id':'huakang-d','expected_revision':doc['revision']}).status_code==422


def test_link_revocation_during_render_blocks_the_prepared_pdf(client,monkeypatch):
    party,headers=issuer(client);doc=confirmed(client,headers)
    files=importlib.import_module('app.services.material_handover_files')
    def revoke(*args):
        database=importlib.import_module('app.db');model=importlib.import_module('app.models.material_handover').HandoverIssuer
        with database.SessionLocal() as db:
            row=db.get(model,party['id']);row.active=0;db.commit()
        return b'%PDF-1.4\nprivate-copy'
    monkeypatch.setattr(files,'render_original',revoke)
    response=client.get(BASE+'/external/documents/'+doc['id']+'/print',headers=headers,params={'expected_revision':doc['revision']})
    assert response.status_code==401 and b'private-copy' not in response.content
