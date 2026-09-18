from io import BytesIO

from openpyxl import load_workbook

from test_molding_sample_api import login_as, make_client
from test_qc_inspection_api import _order_payload


def test_schedule_fallback_sort_report_scope_and_shipment_changes(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "qc_inspector")
        created = {}
        for name, planned, shipment in [
            ("fallback", "", "2026-12-28"),
            ("explicit", "2027-01-04", "2026-12-28"),
            ("missing", "", ""),
            ("early", "2026-12-21", "2027-01-04"),
        ]:
            payload = {**_order_payload(f"date-{name}"), "customer_item_no": name,
                       "planned_inspection_date": planned, "shipment_date": shipment}
            response = client.post("/api/qc-inspections/orders", json=payload)
            assert response.status_code == 201, response.text
            created[name] = response.json()
        rows = client.get("/api/qc-inspections/orders", params={"factory_id": "huaxing"}).json()["items"]
        assert [row["customer_item_no"] for row in rows] == ["early", "fallback", "explicit", "missing"]
        # Raw customer dates stay blank; fallback is derived, never written over source values.
        assert created["fallback"]["planned_inspection_date"] == ""

        def report_items(request_id):
            response = client.post("/api/qc-inspections/reports/generate", json={
                "factory_id": "huaxing", "week_key": "2026-W53", "period_mode": "WEEK",
                "period_key": "2026-W53", "report_type": "WEEKLY_INSPECTION_SCHEDULE", "request_id": request_id,
            })
            assert response.status_code == 201, response.text
            download = client.get(f"/api/qc-inspections/reports/{response.json()['id']}/download", params={"factory_id": "huaxing"})
            assert download.status_code == 200
            workbook = load_workbook(BytesIO(download.content), data_only=True)
            items = list(workbook.active.iter_rows(min_row=3, values_only=True))
            workbook.close()
            return items

        report = report_items("fallback-week")
        assert len(report) == 1
        assert report[0][0] == "2026-12-28"
        assert report[0][4] == "fallback"
        changed = client.patch(f"/api/qc-inspections/orders/{created['fallback']['id']}", json={
            "factory_id": "huaxing", "expected_revision": created["fallback"]["revision"],
            "request_id": "shipment-moved", "reason": "走货期调整", "shipment_date": "2027-01-11",
        })
        assert changed.status_code == 200, changed.text
        assert changed.json()["planned_inspection_date"] == ""
        # Moving shipment automatically moves the effective schedule out of the original reporting week.
        assert not any(row[4] == "fallback" for row in report_items("fallback-moved-week"))
