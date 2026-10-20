from dataclasses import replace
from io import BytesIO

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from app.api.spreadsheet_merge import router
from app.core.config import settings
from app.services.auth import AuthContext, AuthGrantContext, AuthOverrideContext, get_current_user

BASE = '/api/tools/spreadsheet-merge'
PERMISSIONS = frozenset({'carton_procurement:read', 'carton_procurement:inventory_write'})


def user(*, permissions=PERMISSIONS, factory='huaxing', department='pmc-warehouse', **grant_fields):
    grant = AuthGrantContext(role_id='warehouse_keeper', role_name='仓管',
                             factory_id=factory, department=department,
                             permissions=permissions, **grant_fields)
    return AuthContext(id='test', username='test', display_name='测试', roles=('warehouse_keeper',),
                       role_codes=(), permissions=permissions, factory_scopes=(factory,),
                       department_scopes=(department,), grants=(grant,))


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(settings, 'document_tools_enabled', True)
    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_current_user] = lambda: user()
    with TestClient(app) as instance:
        yield instance


def acting(client, context):
    client.app.dependency_overrides[get_current_user] = lambda: context


def catalog(client, factory='huaxing'):
    return client.get(f'{BASE}/templates', params={'factory_id': factory})


def payload(template):
    return {
        'primary': {'name': template['label'], 'headers': template['headers'], 'rows': []},
        'secondary': {'name': 'own.csv', 'headers': ['自有列'], 'rows': []},
        'template_id': template['id'], 'profile': template['profile'], 'factory_id': 'huaxing',
    }


def test_warehouse_catalog_and_each_download_is_factory_scoped(client):
    response = catalog(client)
    assert response.status_code == 200
    assert response.headers['cache-control'] == 'no-store'
    entries = response.json()
    assert {entry['id'] for entry in entries} == {
        'carton-packaging', 'carton-paper-options', 'carton-locations',
        'carton-history-orders', 'carton-opening-inventory', 'carton-weekly-schedule',
    }
    for entry in entries:
        response = client.get(f"{BASE}/templates/{entry['id']}", params={'factory_id': 'huaxing'})
        assert response.status_code == 200
        assert response.headers['cache-control'] == 'no-store'
        book = load_workbook(BytesIO(response.content), read_only=True)
        assert book.sheetnames
        book.close()
        assert client.post(f'{BASE}/preview', json=payload(entry)).status_code == 200
    assert catalog(client, 'huakang-a').json() == []
    assert client.get(f'{BASE}/templates/carton-packaging?factory_id=huakang-a').status_code == 403
    assert catalog(client, '*').status_code == 422
    assert catalog(client, 'not-a-factory').status_code == 422
    assert client.get(f'{BASE}/templates').status_code == 422
    assert client.get(f'{BASE}/templates/unknown?factory_id=huaxing').status_code == 404


@pytest.mark.parametrize('context', [
    user(permissions=frozenset()),
    user(permissions=frozenset({'carton_procurement:read'})),
    user(department='sales-business'),
    user(factory='huakang-a'),
    replace(user(), grants=()),  # Flat permission union cannot authorize.
    replace(user(), account_available=False),
    user(valid_until='2000-01-01T00:00:00+00:00'),
    user(factory='*', factory_ceiling=('huakang-a',)),
])
def test_nonwarehouse_or_inactive_context_has_no_system_templates(client, context):
    entry = catalog(client).json()[0]
    acting(client, context)
    assert catalog(client).json() == []
    assert client.get(f"{BASE}/templates/{entry['id']}?factory_id=huaxing").status_code == 403
    for operation in ('preview', 'export'):
        assert client.post(f'{BASE}/{operation}', json=payload(entry)).status_code == 403
    # Ordinary own-file conversion remains available to authenticated users;
    # actual disabled-account sessions are rejected by get_current_user.
    ordinary = payload(entry)
    ordinary.pop('template_id')
    ordinary.pop('factory_id')
    ordinary['profile'] = 'generic'
    assert client.post(f'{BASE}/preview', json=ordinary).status_code == 200


@pytest.mark.parametrize('mode', ['legacy', 'shadow', 'enforce'])
def test_explicit_deny_is_enforced_in_all_rollout_modes(client, monkeypatch, mode):
    monkeypatch.setattr(settings, 'authz_mode', mode)
    denied = AuthOverrideContext(id='deny', permission_code='carton_procurement:inventory_write',
                                 effect='deny', factory_id='huaxing', department='*')
    acting(client, replace(user(), overrides=(denied,)))
    assert catalog(client).json() == []
    assert client.get(f'{BASE}/templates/carton-packaging?factory_id=huaxing').status_code == 403


def test_admin_scope_and_deny_remain_canonical(client):
    admin = user(factory='*', department='system', role_code='admin', permissions=frozenset())
    acting(client, admin)
    assert len(catalog(client).json()) == 6
    denied = AuthOverrideContext(id='deny', permission_code='carton_procurement:read',
                                 effect='deny', factory_id='huaxing', department='*')
    acting(client, replace(admin, overrides=(denied,)))
    assert catalog(client).json() == []
    assert len(catalog(client, 'huakang-a').json()) == 6


def test_cached_template_and_legacy_profile_cannot_bypass_revocation(client):
    entry = catalog(client).json()[0]
    request = payload(entry)
    assert client.post(f'{BASE}/preview', json=request).status_code == 200
    acting(client, user(permissions=frozenset()))
    for operation in ('preview', 'export'):
        assert client.post(f'{BASE}/{operation}', json=request).status_code == 403
        legacy = {**request, 'template_id': None, 'profile': 'carton-packaging'}
        assert client.post(f'{BASE}/{operation}', json=legacy).status_code == 403
        no_factory = {**request, 'factory_id': None}
        assert client.post(f'{BASE}/{operation}', json=no_factory).status_code == 422
