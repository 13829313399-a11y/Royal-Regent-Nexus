import json
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.services.internal_quote import _rr2_cost_summary
from app.services.internal_quote_excel import _build_summary_sheet
from test_internal_quote_summary import SNAPSHOT, rows_by_key, section, summary_sections
from test_internal_quote_purchase_export import arithmetic_value


@pytest.mark.parametrize('factory', ['huaxing', 'huadeng', 'huakang-a', 'huakang-b', 'huakang-c', 'huakang-d'])
@pytest.mark.parametrize('standalone_hair', [False, True])
def test_material_clothes_and_hair_never_increase_tax_deduction(factory, standalone_hair):
    def summary(extra):
        sections = [s for s in summary_sections(freight_enabled=False)
                    if s.department not in {'molding', 'sewing'}]
        sales = next(s for s in sections if s.department == 'sales')
        sales.payload_json = json.dumps({**json.loads(sales.payload_json), 'shipping': {'misc_ratio': 0}})
        sections += [
            section('molding', {}, {'injection_material_hkd': str(5 + extra),
                    'injection_domestic_material_hkd': str(5 + extra),
                    'injection_labor_hkd': '3', 'total_hkd': str(8 + extra)}),
            section('sewing', {}, {'clothes_material_hkd': str(2 + extra),
                    'clothes_hkd': str(4 + extra), 'hair_hkd': '0' if standalone_hair else str(6 + extra),
                    'total_hkd': str(4 + extra) if standalone_hair else str(10 + extra * 2)}),
        ]
        if standalone_hair:
            sections.append(section('hair', {}, {'total_hkd': str(6 + extra)}))
        return _rr2_cost_summary(sections, {'factory_price_hkd': Decimal(50 + extra * 3)},
                                 SNAPSHOT, factory_id=factory)

    before, after = summary(0), summary(10)
    for key in ('material', 'sewing_cloth', 'sewing_hair'):
        assert Decimal(rows_by_key(after['t1'])[key]['value']) - Decimal(rows_by_key(before['t1'])[key]['value']) == 10
    for result in (before, after):
        taxes = rows_by_key(result['t4'])
        for key in ('sewhair13', 'sewcloth13'):
            assert taxes[key]['rate_percent'] is None
            assert taxes[key]['deduction_hkd'] is None
        assert Decimal(taxes['tax13b']['deduction_hkd']) > 0
        assert taxes['slush3']['rate_percent'] == '3.0000'
    assert rows_by_key(before['t4'])['tax13']['amount_hkd'] == rows_by_key(after['t4'])['tax13']['amount_hkd']
    assert before['totals']['total_deduction_hkd'] == after['totals']['total_deduction_hkd']
    assert Decimal(after['totals']['after_deduction_cost_hkd']) - Decimal(before['totals']['after_deduction_cost_hkd']) == 30


@pytest.mark.parametrize('mode', ['ordinary', 'detached', 'component'])
@pytest.mark.parametrize('hair_section', ['hair', 'sewing'])
def test_export_excludes_three_categories_even_with_old_tax_snapshot(mode, hair_section):
    entries = [
        {'section': 'sewing', 'category': 'clothes', 'label': '公仔布套', 'amount_hkd': 4},
        {'section': hair_section, 'category': 'hair', 'label': '公仔头发', 'amount_hkd': 6},
        {'section': 'engineering', 'category': 'hardware', 'label': '螺丝', 'amount_hkd': 2},
    ]
    shipping = {'markup': 1.2, 'misc_ratio': .02, 'rows': [], 'pricing_entries': entries}
    if mode == 'component':
        for entry in entries:
            entry['pricing_component_id'] = 'a'
        shipping.update(pricing_mode='component', pricing_groups=[{'id': 'a', 'name': '公仔', 'cost_hkd': 20, 'markup': 1.2}],
                        global_pricing={'entries': [], 'cost_hkd': 0, 'markup': 1.2})
    elif mode == 'detached':
        for entry in entries[:2]:
            entry['markup_override'] = 1.03
        shipping.update(pricing_mode='standard', pricing_groups=[
            {'id': 'main', 'name': '主倍率汇总', 'cost_hkd': 10, 'markup': 1.2},
            *[{'id': f'detail-{i}', 'name': entry['label'], 'cost_hkd': entry['amount_hkd'], 'markup': 1.03}
              for i, entry in enumerate(entries[:2])]])
    summary = {'t1': [{'key': k, 'value': v} for k, v in [('material', 5), ('sewing_cloth', 4), ('sewing_hair', 6), ('hardware', 2)]],
               't2': [], 't3': [{'key': 'injection_labor', 'value': 3}],
               't4': [{'key': k, 'label': k, 'amount_hkd': 10, 'rate_percent': 11.5, 'deduction_hkd': 1.15}
                      for k in ['sewhair13', 'sewcloth13', 'tax13b']], 'shipping_pricing': shipping}
    sections = [section('sales', {'pricing_mode': 'component' if mode == 'component' else 'standard'}, {}),
                section('molding', {}, {}, [{'kind': 'injection', 'item': '胶壳', 'material_cost_hkd': 5,
                                             'molding_cost_hkd': 3, 'pricing_component_id': 'a'}])]
    original = json.dumps(summary)
    workbook = Workbook()
    try:
        _build_summary_sheet(workbook, SimpleNamespace(product_name='测试公仔', quote_no='TAX-QA',
                             customer='JustPlay' if mode == 'component' else '普通客', region_code='mainland', remark=''),
                             sections, SNAPSHOT, summary, {'molding_hkd': 8, 'sewing_hkd': 10, 'factory_price_hkd': 20}, [])
        sheet = workbook['报价明细']
        for category, amount in [('料价', 5), ('车衣', 4), ('车发', 6)]:
            rows = [r for r in range(1, sheet.max_row + 1) if sheet.cell(r, 2).value == category]
            assert rows, category
            assert all(sheet.cell(r, 1).value in ('', None) for r in rows)
            assert sum(arithmetic_value(sheet, f'D{r}') for r in rows) == pytest.approx(amount)
        assert any(sheet.cell(r, 1).value == '¥13%' for r in range(1, sheet.max_row + 1)
                   if sheet.cell(r, 2).value == '五金')
        tax_header = next(cell.row for row in sheet for cell in row if cell.value == '人民币外购件成本')
        for column in (8, 9):
            assert sheet.cell(tax_header + 1, column).value in ('', None)
            assert sheet.cell(tax_header + 3, column).value in ('', None)
        assert json.dumps(summary) == original
    finally:
        workbook.close()
