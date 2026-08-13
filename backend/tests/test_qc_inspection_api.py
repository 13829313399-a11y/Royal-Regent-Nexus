import json
import hashlib
import sqlite3
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image
from pypdf import PdfWriter

from test_molding_sample_api import (
    grant_permission_override,
    login_as,
    make_client,
)


def _order_payload(request_id: str = "order-1") -> dict[str, object]:
    return {
        "factory_id": "huaxing",
        "week_key": "2026-W33",
        "customer_name": "彩星",
        "sales_contract_no": "SC001",
        "customer_item_no": "ITEM001",
        "customer_po_no": "001234",
        "product_name": "测试产品",
        "quantity": "1200",
        "packing": "12件/箱",
        "carton_count": "100",
        "report_status": "已完成",
        "production_department": "华兴生产部",
        "export_country_code": "US",
        "shipment_date": "2026-08-20",
        "planned_inspection_date": "2026-08-12",
        "inspection_agency": "QC",
        "account_manager": "香港跟客A",
        "request_id": request_id,
    }


def _schedule_bytes(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "客户",
            "合同号",
            "客户货号",
            "PO",
            "产品名称",
            "数量",
            "出口国",
            "走货日期",
            "计划验货日期",
        ]
    )
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _jpeg_bytes(color: str = "red") -> bytes:
    output = BytesIO()
    Image.new("RGB", (8, 8), color).save(output, format="JPEG")
    return output.getvalue()


def test_qc_order_po_idempotency_pass_manual_problem_and_factory_isolation(monkeypatch):
    with make_client(monkeypatch) as client:
        profile = login_as(client, "qc_inspector")
        assert "qc_inspection:order_write" in profile["permissions"]
        assert "qc_inspection:customer_manage" not in profile["permissions"]

        missing_po = _order_payload("missing-po")
        missing_po["customer_po_no"] = ""
        assert client.post("/api/qc-inspections/orders", json=missing_po).status_code == 422

        created_response = client.post("/api/qc-inspections/orders", json=_order_payload())
        assert created_response.status_code == 201, created_response.text
        created = created_response.json()
        assert created["customer_po_no"] == "001234"
        assert created["inspection_no"].startswith("QC-")

        replay = client.post("/api/qc-inspections/orders", json=_order_payload())
        assert replay.status_code == 201
        assert replay.json()["id"] == created["id"]
        collision = _order_payload()
        collision["quantity"] = "999"
        assert client.post("/api/qc-inspections/orders", json=collision).status_code == 409

        cross_factory = client.get(
            "/api/qc-inspections/orders",
            params={"factory_id": "huakang-a"},
        )
        assert cross_factory.status_code == 403

        updated = client.patch(
            f"/api/qc-inspections/orders/{created['id']}",
            json={
                "factory_id": "huaxing",
                "expected_revision": created["revision"],
                "reason": "现场发现问题",
                "actual_inspection_date": "2026-08-12",
                "inspection_result": "PASS",
                "manual_has_problem": True,
                "request_id": "result-pass-manual",
            },
        )
        assert updated.status_code == 200, updated.text
        assert updated.json()["inspection_result"] == "PASS"
        assert updated.json()["problem_count"] == 1
        assert updated.json()["has_problem"] is True

        problems = client.get(
            "/api/qc-inspections/problems",
            params={"factory_id": "huaxing", "week": "2026-W33"},
        ).json()["items"]
        assert len(problems) == 1
        assert problems[0]["source_type"] == "AUTO_RESULT"
        assert problems[0]["status"] == "DRAFT"

        unchanged_trigger = client.patch(
            f"/api/qc-inspections/orders/{created['id']}",
            json={
                "factory_id": "huaxing",
                "expected_revision": updated.json()["revision"],
                "reason": "补充机构",
                "inspection_agency": "第三方机构",
                "request_id": "result-pass-manual-2",
            },
        )
        assert unchanged_trigger.status_code == 200
        assert unchanged_trigger.json()["problem_count"] == 1

        cancelled = client.post(
            "/api/qc-inspections/orders", json=_order_payload("cancelled-order")
        ).json()
        cancelled_result = client.patch(
            f"/api/qc-inspections/orders/{cancelled['id']}",
            json={
                "factory_id": "huaxing",
                "expected_revision": cancelled["revision"],
                "reason": "客户取消验货",
                "actual_inspection_date": "2026-08-12",
                "inspection_result": "CANCELLED",
                "request_id": "cancelled-result",
            },
        )
        assert cancelled_result.status_code == 200
        assert cancelled_result.json()["problem_count"] == 1
        reopened = client.patch(
            f"/api/qc-inspections/orders/{cancelled['id']}",
            json={
                "factory_id": "huaxing",
                "expected_revision": cancelled_result.json()["revision"],
                "reason": "撤销错误取消结果",
                "inspection_result": "PENDING",
                "request_id": "cancelled-back-to-pending",
            },
        )
        assert reopened.status_code == 200
        assert reopened.json()["status"] == "SCHEDULED"

        no_actual_date = client.post(
            "/api/qc-inspections/orders", json=_order_payload("no-actual-date")
        ).json()
        assert client.patch(
            f"/api/qc-inspections/orders/{no_actual_date['id']}",
            json={
                "factory_id": "huaxing",
                "expected_revision": no_actual_date["revision"],
                "reason": "结果缺日期",
                "inspection_result": "FAIL",
                "request_id": "no-actual-date-result",
            },
        ).status_code == 422


def test_problem_fields_are_maintained_only_on_problem_resource(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "qc_inspector")
        order = client.post(
            "/api/qc-inspections/orders", json=_order_payload("problem-owner")
        ).json()
        illegal_order_write = client.patch(
            f"/api/qc-inspections/orders/{order['id']}",
            json={
                "factory_id": "huaxing",
                "expected_revision": order["revision"],
                "reason": "错误入口",
                "problem_summary": "不能在主单直接维护",
                "request_id": "problem-wrong-resource",
            },
        )
        assert illegal_order_write.status_code == 422

        problem = client.post(
            "/api/qc-inspections/problems",
            json={
                "factory_id": "huaxing",
                "inspection_order_id": order["id"],
                "status": "OPEN",
                "category": "包装",
                "description": "外箱破损",
                "return_reason": "包装不合格",
                "corrective_action": "更换外箱",
                "resolution": "",
                "reported_date": "2026-08-12",
                "request_id": "problem-create",
            },
        )
        assert problem.status_code == 201, problem.text
        detail = client.get(
            f"/api/qc-inspections/orders/{order['id']}",
            params={"factory_id": "huaxing"},
        )
        assert detail.status_code == 200
        assert detail.json()["problem_summary"] == "外箱破损"


def test_schedule_preview_multiple_candidate_requires_explicit_target(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "qc_inspector")
        first = client.post("/api/qc-inspections/orders", json=_order_payload("multi-1")).json()
        second_payload = _order_payload("multi-2")
        second_payload["customer_po_no"] = "001235"
        second = client.post("/api/qc-inspections/orders", json=second_payload).json()
        content = _schedule_bytes(
            [["彩星", "SC001", "ITEM001", "001236", "新产品", 1200, "US", "2026-08-21", "2026-08-13"]]
        )
        preview = client.post(
            "/api/qc-inspections/schedule-imports/preview",
            data={
                "factory_id": "huaxing",
                "week_key": "2026-W33",
                "request_id": "schedule-preview-1",
            },
            files={"file": ("排期.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert preview.status_code == 201, preview.text
        batch = preview.json()
        row = batch["rows"][0]
        assert row["match_status"] == "MULTIPLE_MATCHES"
        assert set(row["candidate_order_ids"]) == {first["id"], second["id"]}

        missing_target = client.post(
            f"/api/qc-inspections/schedule-imports/{batch['id']}/confirm",
            json={
                "factory_id": "huaxing",
                "expected_revision": batch["revision"],
                "request_id": "schedule-confirm-missing-target",
                "decisions": [{"row_id": row["id"], "action": "UPDATE"}],
            },
        )
        assert missing_target.status_code == 422

        confirmed = client.post(
            f"/api/qc-inspections/schedule-imports/{batch['id']}/confirm",
            json={
                "factory_id": "huaxing",
                "expected_revision": batch["revision"],
                "request_id": "schedule-confirm-target",
                "decisions": [
                    {"row_id": row["id"], "action": "UPDATE", "target_order_id": first["id"]}
                ],
            },
        )
        assert confirmed.status_code == 200, confirmed.text
        assert confirmed.json()["status"] == "CONFIRMED"


def test_schedule_history_is_not_candidate_and_stale_candidate_is_rejected(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "qc_inspector")
        historical = client.post(
            "/api/qc-inspections/orders", json=_order_payload("historical-order")
        ).json()
        completed = client.patch(
            f"/api/qc-inspections/orders/{historical['id']}",
            json={
                "factory_id": "huaxing",
                "expected_revision": historical["revision"],
                "reason": "历史验货完成",
                "actual_inspection_date": "2026-08-12",
                "inspection_result": "PASS",
                "request_id": "historical-result",
            },
        )
        assert completed.status_code == 200
        content = _schedule_bytes(
            [["彩星", "SC001", "ITEM001", "001999", "新一轮", 100, "US", "2026-08-22", "2026-08-14"]]
        )
        history_preview = client.post(
            "/api/qc-inspections/schedule-imports/preview",
            data={"factory_id": "huaxing", "week_key": "2026-W33", "request_id": "history-preview"},
            files={"file": ("排期.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).json()
        assert history_preview["rows"][0]["match_status"] == "NEW"
        assert history_preview["rows"][0]["candidate_order_ids"] == []

        open_order = client.post(
            "/api/qc-inspections/orders", json=_order_payload("stale-order")
        ).json()
        stale_content = _schedule_bytes(
            [["彩星", "SC001", "ITEM001", "001888", "更新品", 200, "US", "2026-08-23", "2026-08-15"]]
        )
        stale_preview = client.post(
            "/api/qc-inspections/schedule-imports/preview",
            data={"factory_id": "huaxing", "week_key": "2026-W33", "request_id": "stale-preview"},
            files={"file": ("排期.xlsx", stale_content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).json()
        row = stale_preview["rows"][0]
        assert row["candidate_order_revisions"][open_order["id"]] == open_order["revision"]
        assert row["candidate_orders"][0]["inspection_no"] == open_order["inspection_no"]
        changed = client.patch(
            f"/api/qc-inspections/orders/{open_order['id']}",
            json={
                "factory_id": "huaxing",
                "expected_revision": open_order["revision"],
                "reason": "预览后人工修改",
                "note": "已变化",
                "request_id": "stale-order-change",
            },
        )
        assert changed.status_code == 200
        confirm = client.post(
            f"/api/qc-inspections/schedule-imports/{stale_preview['id']}/confirm",
            json={
                "factory_id": "huaxing",
                "expected_revision": stale_preview["revision"],
                "request_id": "stale-confirm",
                "decisions": [{"row_id": row["id"], "action": "UPDATE", "target_order_id": open_order["id"]}],
            },
        )
        assert confirm.status_code == 409


def test_schedule_preview_raw_form_validation(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "qc_inspector")
        content = _schedule_bytes(
            [["彩星", "SC1", "I1", "0001", "产品", 1, "US", "2026-08-20", "2026-08-12"]]
        )
        for week_key, request_id in (("2026-W54", "valid-id"), ("2026-W33", "bad id")):
            response = client.post(
                "/api/qc-inspections/schedule-imports/preview",
                data={"factory_id": "huaxing", "week_key": week_key, "request_id": request_id},
                files={"file": ("排期.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
            assert response.status_code == 422
        import app.api.qc_inspection as qc_api

        monkeypatch.setattr(qc_api, "MAX_SCHEDULE_IMPORT_BYTES", 16)
        oversized = client.post(
            "/api/qc-inspections/schedule-imports/preview",
            data={"factory_id": "huaxing", "week_key": "2026-W33", "request_id": "schedule-too-large"},
            files={"file": ("排期.xlsx", b"x" * 17, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert oversized.status_code == 413


def test_schedule_all_unchanged_preview_is_immediately_confirmed(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "qc_inspector")
        payload = _order_payload("unchanged-order")
        payload.update(
            packing="",
            carton_count="",
            report_status="",
            production_department="",
            inspection_agency="",
            account_manager="",
        )
        order = client.post("/api/qc-inspections/orders", json=payload)
        assert order.status_code == 201
        content = _schedule_bytes(
            [["彩星", "SC001", "ITEM001", "001234", "测试产品", 1200, "US", "2026-08-20", "2026-08-12"]]
        )
        preview = client.post(
            "/api/qc-inspections/schedule-imports/preview",
            data={"factory_id": "huaxing", "week_key": "2026-W33", "request_id": "unchanged-preview"},
            files={"file": ("排期.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert preview.status_code == 201, preview.text
        assert preview.json()["status"] == "CONFIRMED"
        assert preview.json()["rows"][0]["decision_status"] == "UNCHANGED"


def test_rename_preview_execute_persists_sources_and_recovers_same_zip(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "qc_inspector")
        pdf = _pdf_bytes()
        good_jpg = _jpeg_bytes("red")
        bad_pdf = _pdf_bytes()
        metadata = {
            "factory_id": "huaxing",
            "request_id": "rename-preview-1",
            "groups": [
                {
                    "group_id": "good",
                    "is_caixing": False,
                    "item_number": "ITEM001",
                    "customer_po_no": "001234",
                    "actual_inspection_date": "2026-08-12",
                    "export_country": "US",
                    "files": [
                        {"file_id": "good-pdf", "source_file_name": "报告.pdf"},
                        {"file_id": "good-jpg", "source_file_name": "照片.jpg"},
                    ],
                },
                {
                    "group_id": "blocked",
                    "is_caixing": False,
                    "item_number": "ITEM002",
                    "customer_po_no": "000002",
                    "actual_inspection_date": "2026-08-12",
                    "export_country": "DE",
                    "files": [
                        {"file_id": "bad-pdf", "source_file_name": "仅报告.pdf"},
                    ],
                },
            ],
        }
        data = {
            "metadata_json": json.dumps(metadata, ensure_ascii=False),
            "file_ids": ["good-pdf", "good-jpg", "bad-pdf"],
        }
        files = [
            ("files", ("报告.pdf", pdf, "application/pdf")),
            ("files", ("照片.jpg", good_jpg, "image/jpeg")),
            ("files", ("仅报告.pdf", bad_pdf, "application/pdf")),
        ]
        preview = client.post(
            "/api/qc-inspections/rename-batches/preview", data=data, files=files
        )
        assert preview.status_code == 201, preview.text
        batch = preview.json()
        assert batch["successful_group_count"] == 1
        assert batch["failed_group_count"] == 1
        good = next(item for item in batch["groups"] if item["group_id"] == "good")
        assert {item["target_file_name"] for item in good["files"]} == {
            "US_ITEM001_001234_20260812.pdf",
            "US_ITEM001_001234_20260812.jpg",
        }

        replay = client.post(
            "/api/qc-inspections/rename-batches/preview", data=data, files=files
        )
        assert replay.status_code == 201
        assert replay.json()["id"] == batch["id"]

        db_module = __import__("app.db", fromlist=["SessionLocal"])
        models = __import__("app.models.qc_inspection", fromlist=["QcReportRenameSourceFile"])
        with db_module.SessionLocal() as db:
            sources = db.query(models.QcReportRenameSourceFile).all()
            source_by_name = {item.source_file_name: bytes(item.source_bytes) for item in sources}
            assert source_by_name["报告.pdf"] == pdf
            assert source_by_name["照片.jpg"] == good_jpg
            assert source_by_name["仅报告.pdf"] == bad_pdf

        executed = client.post(
            f"/api/qc-inspections/rename-batches/{batch['id']}/execute",
            json={
                "factory_id": "huaxing",
                "expected_revision": batch["revision"],
                "request_id": "rename-execute-1",
            },
        )
        assert executed.status_code == 200, executed.text
        first_sha = hashlib.sha256(executed.content).hexdigest()
        with ZipFile(BytesIO(executed.content)) as archive:
            assert sorted(archive.namelist()) == [
                "US_ITEM001_001234_20260812.jpg",
                "US_ITEM001_001234_20260812.pdf",
            ]
            assert archive.read("US_ITEM001_001234_20260812.pdf") == pdf

        recovered = client.post(
            f"/api/qc-inspections/rename-batches/{batch['id']}/execute",
            json={
                "factory_id": "huaxing",
                "expected_revision": batch["revision"],
                "request_id": "rename-execute-recovery",
            },
        )
        assert recovered.status_code == 200
        assert hashlib.sha256(recovered.content).hexdigest() == first_sha
        assert recovered.content == executed.content

        with db_module.SessionLocal() as db:
            stored_batch = db.get(models.QcReportRenameBatch, batch["id"])
            assert stored_batch is not None
            stored_batch.archive_size_bytes = len(executed.content) + 1
            db.commit()

        corrupt_size_replay = client.post(
            f"/api/qc-inspections/rename-batches/{batch['id']}/execute",
            json={
                "factory_id": "huaxing",
                "expected_revision": batch["revision"],
                "request_id": "rename-execute-1",
            },
        )
        assert corrupt_size_replay.status_code == 409
        assert "完整性" in corrupt_size_replay.json()["detail"]

        with db_module.SessionLocal() as db:
            stored_batch = db.get(models.QcReportRenameBatch, batch["id"])
            assert stored_batch is not None
            stored_batch.archive_size_bytes = len(executed.content)
            stored_batch.archive_sha256 = "0" * 64
            db.commit()

        corrupt_hash_recovery = client.post(
            f"/api/qc-inspections/rename-batches/{batch['id']}/execute",
            json={
                "factory_id": "huaxing",
                "expected_revision": batch["revision"],
                "request_id": "rename-execute-corrupt-recovery",
            },
        )
        assert corrupt_hash_recovery.status_code == 409
        assert "完整性" in corrupt_hash_recovery.json()["detail"]

        corrupt_download = client.get(
            f"/api/qc-inspections/rename-batches/{batch['id']}/download",
            params={"factory_id": "huaxing"},
        )
        assert corrupt_download.status_code == 409
        assert "完整性" in corrupt_download.json()["detail"]

        cross_factory = client.post(
            f"/api/qc-inspections/rename-batches/{batch['id']}/execute",
            json={
                "factory_id": "huakang-a",
                "expected_revision": 2,
                "request_id": "rename-cross-factory",
            },
        )
        assert cross_factory.status_code == 403


def test_rename_preview_rejects_casefold_duplicate_names_and_oversized_file(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "qc_inspector")
        duplicate = {
            "factory_id": "huaxing",
            "request_id": "rename-duplicate-name",
            "groups": [{
                "group_id": "dup",
                "is_caixing": False,
                "item_number": "I1",
                "customer_po_no": "0001",
                "actual_inspection_date": "2026-08-12",
                "export_country": "US",
                "files": [
                    {"file_id": "one", "source_file_name": "Same.PDF"},
                    {"file_id": "two", "source_file_name": "same.pdf"},
                ],
            }],
        }
        response = client.post(
            "/api/qc-inspections/rename-batches/preview",
            data={
                "metadata_json": json.dumps(duplicate),
                "file_ids": ["one", "two"],
            },
            files=[
                ("files", ("Same.PDF", _pdf_bytes(), "application/pdf")),
                ("files", ("same.pdf", _pdf_bytes(), "application/pdf")),
            ],
        )
        assert response.status_code == 422

        import app.api.qc_inspection as qc_api

        monkeypatch.setattr(qc_api, "MAX_FILE_SIZE_BYTES", 16)
        oversized = {
            "factory_id": "huaxing",
            "request_id": "rename-too-large",
            "groups": [{
                "group_id": "large",
                "is_caixing": False,
                "item_number": "I1",
                "customer_po_no": "0001",
                "actual_inspection_date": "2026-08-12",
                "export_country": "US",
                "files": [{"file_id": "large-file", "source_file_name": "large.pdf"}],
            }],
        }
        too_large = client.post(
            "/api/qc-inspections/rename-batches/preview",
            data={"metadata_json": json.dumps(oversized), "file_ids": ["large-file"]},
            files=[("files", ("large.pdf", b"x" * 17, "application/pdf"))],
        )
        assert too_large.status_code == 413


def test_sqlite_foreign_keys_and_two_session_revision_conflict(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "qc_inspector")
        created = client.post(
            "/api/qc-inspections/orders", json=_order_payload("concurrent-order")
        ).json()
        db_module = __import__("app.db", fromlist=["SessionLocal"])
        model_module = __import__("app.models.qc_inspection", fromlist=["QcInspectionOrder"])
        schema_module = __import__("app.schemas.qc_inspection", fromlist=["QcOrderUpdate"])
        service_module = __import__("app.services.qc_inspection", fromlist=["update_order"])
        auth_module = __import__("app.services.auth", fromlist=["AuthContext"])
        with db_module.engine.connect() as connection:
            assert connection.exec_driver_sql("PRAGMA foreign_keys").scalar_one() == 1

        session_one = db_module.SessionLocal()
        session_two = db_module.SessionLocal()
        try:
            assert session_one.get(model_module.QcInspectionOrder, created["id"]).revision == 1
            assert session_two.get(model_module.QcInspectionOrder, created["id"]).revision == 1
            actor = auth_module.AuthContext(
                id="user-qc-inspector",
                username="qc_inspector",
                display_name="华兴QC检验员",
                roles=("QC检验员",),
                role_codes=("position_qc_inspector",),
                permissions=frozenset(),
                factory_scopes=("huaxing",),
                department_scopes=("qc",),
            )
            first = service_module.update_order(
                session_one,
                created["id"],
                schema_module.QcOrderUpdate(
                    factory_id="huaxing",
                    expected_revision=1,
                    reason="并发一",
                    note="first",
                    request_id="concurrent-first",
                ),
                actor,
            )
            assert first["revision"] == 2
            with pytest.raises(Exception) as conflict:
                service_module.update_order(
                    session_two,
                    created["id"],
                    schema_module.QcOrderUpdate(
                        factory_id="huaxing",
                        expected_revision=1,
                        reason="并发二",
                        note="second",
                        request_id="concurrent-second",
                    ),
                    actor,
                )
            assert getattr(conflict.value, "status_code", None) == 409
        finally:
            session_one.close()
            session_two.close()


def test_qc_schema_guard_rejects_existing_unmigrated_alembic_database(monkeypatch):
    with make_client(monkeypatch):
        db_module = __import__("app.db", fromlist=["ensure_qc_inspection_schema_ready"])
        from sqlalchemy import create_engine

        temp_root = Path(".tmp-tests")
        temp_root.mkdir(exist_ok=True)
        database_path = temp_root / f"pre-qc-{uuid4().hex}.sqlite3"
        connection = sqlite3.connect(database_path)
        try:
            connection.execute("CREATE TABLE alembic_version (version_num VARCHAR(32) NOT NULL)")
            connection.execute("INSERT INTO alembic_version VALUES ('20260812_0066')")
            connection.commit()
        finally:
            connection.close()
        guard_engine = create_engine(f"sqlite:///{database_path}", future=True)
        monkeypatch.setattr(db_module, "engine", guard_engine)
        with pytest.raises(RuntimeError, match="20260812_0067"):
            db_module.ensure_qc_inspection_schema_ready()
        guard_engine.dispose()
        database_path.unlink(missing_ok=True)


def test_five_reports_are_real_xlsx_with_expected_layout_and_formal_permission(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "qc_supervisor")
        empty_customer = client.post(
            "/api/qc-inspections/customers",
            json={
                "factory_id": "huaxing",
                "customer_code": "EMPTY",
                "customer_name": "空客户",
                "request_id": "empty-customer-config",
            },
        )
        assert empty_customer.status_code == 201, empty_customer.text
        login_as(client, "qc_inspector")
        order = client.post("/api/qc-inspections/orders", json=_order_payload("report-order")).json()
        result_response = client.patch(
            f"/api/qc-inspections/orders/{order['id']}",
            json={
                "factory_id": "huaxing",
                "expected_revision": order["revision"],
                "reason": "退货",
                "actual_inspection_date": "2026-08-12",
                "inspection_result": "REJECTED",
                "request_id": "report-result",
            },
        )
        assert result_response.status_code == 200
        auto_problem = client.get(
            "/api/qc-inspections/problems",
            params={"factory_id": "huaxing", "week": "2026-W33"},
        ).json()["items"][0]
        problem_update = client.patch(
            f"/api/qc-inspections/problems/{auto_problem['id']}",
            json={
                "factory_id": "huaxing",
                "expected_revision": auto_problem["revision"],
                "reason": "补齐问题责任资料",
                "status": "OPEN",
                "description": "外箱破损",
                "return_reason": "包装不合格",
                "primary_responsible_person": "责任人甲",
                "secondary_responsible_person": "责任人乙",
                "request_id": "report-problem-details",
            },
        )
        assert problem_update.status_code == 200, problem_update.text
        pending = client.post(
            "/api/qc-inspections/orders", json=_order_payload("report-pending")
        )
        assert pending.status_code == 201
        cancelled = client.post(
            "/api/qc-inspections/orders", json=_order_payload("report-cancelled")
        ).json()
        assert client.patch(
            f"/api/qc-inspections/orders/{cancelled['id']}",
            json={
                "factory_id": "huaxing",
                "expected_revision": cancelled["revision"],
                "reason": "取消验货",
                "actual_inspection_date": "2026-08-12",
                "inspection_result": "CANCELLED",
                "request_id": "report-cancelled-result",
            },
        ).status_code == 200
        formal_denied = client.post(
            "/api/qc-inspections/reports/generate",
            json={
                "factory_id": "huaxing",
                "week_key": "2026-W33",
                "report_type": "WEEKLY_STATISTICS",
                "is_formal_snapshot": True,
                "request_id": "formal-denied",
            },
        )
        assert formal_denied.status_code == 403

        expectations = {
            "CUSTOMER_SUMMARY": (18, "彩星验货汇总"),
            "WEEKLY_STATISTICS": (18, "2026年每周客验货情况统计表"),
            "HUAXING_CUSTOMER_WEEKLY_DETAIL": (18, "华兴各客户每周验货明细（2026-W33）"),
            "HUAXING_WEEKLY_AGGREGATE": (8, "华兴每周验货汇总（2026-W33）"),
        }
        customer_source_revision = ""
        for report_type, (column_count, title) in expectations.items():
            generated = client.post(
                "/api/qc-inspections/reports/generate",
                json={
                    "factory_id": "huaxing",
                    "week_key": "2026-W33",
                    "report_type": report_type,
                    "request_id": f"report-{report_type}",
                },
            )
            assert generated.status_code == 201, generated.text
            if report_type == "CUSTOMER_SUMMARY":
                customer_source_revision = generated.json()["source_revision_sha256"]
                assert generated.json()["generation_summary"]["customer_config_count"] == 1
            assert generated.json()["generation_summary"]["order_count"] == 1
            assert generated.json()["generation_summary"]["excluded_pending_count"] == 1
            assert generated.json()["generation_summary"]["excluded_cancelled_count"] == 1
            downloaded = client.get(
                f"/api/qc-inspections/reports/{generated.json()['id']}/download",
                params={"factory_id": "huaxing"},
            )
            assert downloaded.status_code == 200
            workbook = load_workbook(BytesIO(downloaded.content), data_only=False)
            sheet = workbook.worksheets[0]
            assert sheet.max_column == column_count
            assert sheet["A1"].value == title
            assert sheet.page_setup.orientation == "portrait"
            assert sheet.page_setup.fitToWidth == 1
            if report_type in {
                "CUSTOMER_SUMMARY",
                "WEEKLY_STATISTICS",
                "HUAXING_CUSTOMER_WEEKLY_DETAIL",
            }:
                assert sheet.cell(3, 2).value.date().isoformat() == "2026-08-12"
                assert sheet.cell(3, 9).value == "12件/箱"
                assert sheet.cell(3, 10).value == "100"
                assert sheet.cell(3, 12).value == "已完成"
                assert sheet.cell(3, 13).value == "华兴生产部"
                assert sheet.cell(3, 15).value == "外箱破损"
                assert sheet.cell(3, 16).value == "责任人甲"
                assert sheet.cell(3, 17).value == "责任人乙"
            if report_type == "CUSTOMER_SUMMARY":
                assert "空客户" in workbook.sheetnames
                assert workbook["空客户"]["A1"].value == "空客户验货汇总"
                assert workbook["空客户"]["A3"].value == "总计"
            workbook.close()

        db_module = __import__("app.db", fromlist=["SessionLocal"])
        model_module = __import__("app.models.qc_inspection", fromlist=["QcInspectionReport"])
        with db_module.SessionLocal() as db:
            tampered_report = db.get(
                model_module.QcInspectionReport,
                generated.json()["id"],
            )
            tampered_report.artifact_sha256 = "0" * 64
            db.commit()
        tampered_download = client.get(
            f"/api/qc-inspections/reports/{generated.json()['id']}/download",
            params={"factory_id": "huaxing"},
        )
        assert tampered_download.status_code == 409
        assert "完整性" in tampered_download.json()["detail"]

        group_denied = client.post(
            "/api/qc-inspections/reports/generate",
            json={
                "factory_id": "*",
                "week_key": "2026-W33",
                "report_type": "GROUP_SUMMARY",
                "request_id": "group-denied",
            },
        )
        assert group_denied.status_code == 403

        login_as(client, "qc_supervisor")
        config_changed = client.patch(
            f"/api/qc-inspections/customers/{empty_customer.json()['id']}",
            json={
                "factory_id": "huaxing",
                "expected_revision": empty_customer.json()["revision"],
                "customer_name": "空客户（更新）",
                "request_id": "empty-customer-update",
            },
        )
        assert config_changed.status_code == 200
        regenerated_customer = client.post(
            "/api/qc-inspections/reports/generate",
            json={
                "factory_id": "huaxing",
                "week_key": "2026-W33",
                "report_type": "CUSTOMER_SUMMARY",
                "request_id": "customer-after-config-change",
            },
        )
        assert regenerated_customer.status_code == 201
        assert regenerated_customer.json()["source_revision_sha256"] != customer_source_revision
        formal = client.post(
            "/api/qc-inspections/reports/generate",
            json={
                "factory_id": "huaxing",
                "week_key": "2026-W33",
                "report_type": "WEEKLY_STATISTICS",
                "is_formal_snapshot": True,
                "request_id": "formal-allowed",
            },
        )
        assert formal.status_code == 201

        grant_permission_override(
            "qc_supervisor",
            "qc_inspection:group_summary",
            factory_id="*",
            department="*",
        )
        login_as(client, "qc_supervisor")
        group = client.post(
            "/api/qc-inspections/reports/generate",
            json={
                "factory_id": "*",
                "week_key": "2026-W33",
                "report_type": "GROUP_SUMMARY",
                "request_id": "group-authorized",
            },
        )
        assert group.status_code == 201, group.text
        group_download = client.get(
            f"/api/qc-inspections/reports/{group.json()['id']}/download",
            params={"factory_id": "*"},
        )
        assert group_download.status_code == 200
        group_workbook = load_workbook(BytesIO(group_download.content), data_only=False)
        group_sheet = group_workbook.active
        assert group_sheet.max_column == 8
        assert group_sheet["A1"].value == "华登集团验货汇总（2026-W33）"
        assert group_sheet["A3"].value == "华兴"
        assert group_sheet.cell(group_sheet.max_row, 1).value == "总计"
        assert group_sheet.cell(group_sheet.max_row, 5).value == 1
        group_workbook.close()
