"""EDU product-tail and outer-pack interpretation retain source evidence."""
from datetime import date
from io import BytesIO
import json
import os
from pathlib import Path

import openpyxl
import pytest
import xlwt

from app.services.huaxing_order_legacy import edu_po_parser, edu_schedule


@pytest.mark.parametrize("source,expected", [
    ("F12-ZZ901-C9010000", "F12-ZZ901-C901"),
    ("F12-YY902-B9020001", "F12-YY902-B902"),
    ("F12-XX903-C9030021", "F12-XX903-C903"),
    ("F12-WW904-C9040100", "F12-WW904-C904"),
    ("F12-VV905-09010011", "F12-VV905-0901"),
    ("F12-UU906-09020000", "F12-UU906-0902"),
    ("F11-TT907-C9070001", "F11-TT907-C907"),
    ("F12-TT907-C9070001", "F12-TT907-C907"),
    (" F12-ZZ901-C9010000 ", "F12-ZZ901-C901"),
])
def test_evidenced_long_items_remove_only_the_four_digit_tail(source, expected):
    assert edu_schedule.normalize_edu_item(source) == expected
    assert edu_schedule.normalize_edu_item(expected) == expected


@pytest.mark.parametrize("item", [
    "F12-ZZ901-C901", "F12-VV905-0901", "F11-TT907-C907", "F12-TT907-C907",
    "F12-ZZ901-C901001", "F12-ZZ901-C90100001", "F13-ZZ901-C9010000",
    "OTHER-ZZ901-C9010000", "12340001", "ITEM-0001", "F12-ZZ901-C901ABCD",
])
def test_short_or_unproven_item_shapes_are_not_truncated(item):
    assert edu_schedule.normalize_edu_item(item) == item


@pytest.mark.parametrize("raw,expected", [
    ("1/16", 16), ("0/6", 6), ("12/120", 120), (" 6 / 24 ", 24),
    (16, 16), (16.0, 16), ("16", 16), ("1,200", 1200),
    (None, None), ("", None), ("1/0", None), ("1/16/24", None),
    ("待确认1/16", None), ("1/16套", None), (0, None), (-1, None), (True, None),
])
def test_edu_outer_pack_uses_the_right_hand_count_without_guessing(raw, expected):
    row = dict(customer_type="EDU", item_no="F12-ZZ901-C9010000", quantity=120, case_pack=raw)
    edu_schedule.add_derived_fields(row, today=date(2026, 10, 16))
    assert row["case_pack"] == expected
    assert row["source_case_pack"] == raw
    assert row["source_item_no"] == "F12-ZZ901-C9010000"
    assert row["cartons"] == (round(120 / expected, 2) if expected else None)
    edu_schedule.add_derived_fields(row, today=date(2026, 10, 16))
    assert row["item_no"] == "F12-ZZ901-C901"
    assert row["source_item_no"] == "F12-ZZ901-C9010000"
    assert row["source_case_pack"] == raw


def test_non_edu_interpretation_is_unchanged():
    row = dict(customer_type="彩星", item_no="F12-ZZ901-C9010000", quantity=120, case_pack="1/16")
    edu_schedule.add_derived_fields(row)
    assert row["item_no"] == "F12-ZZ901-C9010000"
    assert row["case_pack"] == 1
    assert "source_item_no" not in row and "source_case_pack" not in row


def test_source_item_evidence_retains_original_whitespace():
    original = " \tF12-ZZ901-C9010000\n"
    row = dict(customer_type="EDU", item_no=original)
    edu_schedule.add_derived_fields(row)
    assert row["item_no"] == "F12-ZZ901-C901"
    assert row["source_item_no"] == original


def workbook_bytes(kind, *, edu=True):
    headers = ["客人", "合同", "产品型号", "数量", "装箱", "扣数PO#"]
    values = ["Sample Customer", "SAMPLE-C001", "F12-YY902-B9020000", 160, "1/16", "PO#900000001"]
    out = BytesIO()
    if kind == "xls":
        book = xlwt.Workbook()
        sheet = book.add_sheet("订单")
        if edu:
            book.add_sheet("EDU日期格式")
        for row, items in ((2, headers), (3, values)):
            for col, value in enumerate(items):
                sheet.write(row, col, value)
        book.save(out)
    else:
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "订单"
        if edu:
            book.create_sheet("EDU日期格式")
        for row, items in ((3, headers), (4, values)):
            for col, value in enumerate(items, 1):
                sheet.cell(row, col, value)
        book.save(out)
        book.close()
    return out.getvalue()


@pytest.mark.parametrize("kind", ["xls", "xlsx", "xlsm"])
def test_chinese_edu_reference_sheet_identifies_customer_before_numeric_coercion(kind):
    payload = workbook_bytes(kind)
    parsed = edu_po_parser.parse_po_file(payload, filename=f"华兴代工订单.{kind}")
    assert parsed["customer_type"] == "EDU"
    row, = parsed["rows"]
    assert (row["item_no"], row["case_pack"], row["cartons"]) == ("F12-YY902-B902", 16, 10)
    assert (row["source_item_no"], row["source_case_pack"]) == ("F12-YY902-B9020000", "1/16")
    assert (row["source_sheet"], row["source_row"]) == ("订单", 4)
    assert (row["contract_no"], row["customer_po"], row["quantity"]) == ("SAMPLE-C001", "PO#900000001", 160)


def test_non_edu_workbook_does_not_inherit_customer_specific_tail_rule():
    row, = edu_po_parser.parse_excel_po(workbook_bytes("xlsx", edu=False), filename="彩星订单.xlsx")["rows"]
    assert (row["item_no"], row["case_pack"]) == ("F12-YY902-B9020000", 1)


def test_repeated_derived_field_calls_do_not_override_an_explicit_pack_change():
    row, = edu_po_parser.parse_excel_po(workbook_bytes("xlsx"), filename="华兴订单.xlsx")["rows"]
    row["case_pack"], row["cartons"] = 20, None
    edu_schedule.add_derived_fields(row)
    assert row["case_pack"] == 20 and row["cartons"] == 8
    assert row["source_case_pack"] == "1/16"


def _private_source_case(name):
    """Private file paths and expected business rows stay in an untracked manifest."""
    configured = os.environ.get("HUAXING_EDU_TEST_MANIFEST")
    if not configured:
        pytest.skip("Set HUAXING_EDU_TEST_MANIFEST to a private local JSON manifest")
    manifest_path = Path(configured)
    case = json.loads(manifest_path.read_text(encoding="utf-8"))[name]
    path = Path(case["path"])
    if not path.is_absolute():
        path = manifest_path.parent / path
    return path, case


def test_private_edu_po_keeps_expected_source_rows_and_original_bytes():
    path, case = _private_source_case("po")
    original = path.read_bytes()
    parsed = edu_po_parser.parse_po_file(original, filename=path.name)
    rows = parsed["rows"]
    expected = case["expected_rows"]
    assert expected and len(rows) == len(expected)
    assert all(isinstance(item, dict) and item for item in expected)
    assert [{key: row[key] for key in item} for row, item in zip(rows, expected, strict=True)] == expected
    assert path.read_bytes() == original


def test_private_schedule_codes_and_packing_remain_read_only():
    path, case = _private_source_case("schedule")
    original = path.read_bytes()
    book = openpyxl.load_workbook(BytesIO(original), read_only=True, data_only=True)
    try:
        rows = [(values[7], values[11]) for values in book["ITEM表"].iter_rows(min_row=4, max_col=12, values_only=True)
                if isinstance(values[7], str) and values[7].startswith("F")]
    finally:
        book.close()
    assert len(rows) == case["expected_row_count"]
    assert len({item for item, _ in rows}) == case["expected_unique_product_count"]
    for item, _ in rows:
        normalized = edu_schedule.normalize_edu_item(item)
        assert normalized == item[:-4]
        assert edu_schedule.normalize_edu_item(normalized) == normalized
    for pack, expected_count in case["expected_pack_counts"].items():
        assert sum(str(value) == pack for _, value in rows) == expected_count
    assert path.read_bytes() == original
