import importlib
import sys
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from app.services.injection_scheduling.calculations import (
    effective_changeover_hours,
    estimated_material_kg,
    normalize_ratio,
    parse_machine_a_value,
    progress,
    qualified_shots,
)
from app.services.injection_scheduling.template_contract import UNIFIED_PLAN_HEADERS
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

TEST_TMP_DIR = Path(__file__).resolve().parents[1] / ".pytest-tmp"
ADMIN_TEST_PASSWORD = "AdminSeed123!"


def make_client(monkeypatch) -> tuple[TestClient, object]:
    TEST_TMP_DIR.mkdir(exist_ok=True)
    database_url = f"sqlite:///{TEST_TMP_DIR / f'injection_phase1_{uuid4().hex}.db'}"
    monkeypatch.setenv("DATABASE_URL", database_url)
    monkeypatch.setenv("SEED_ADMIN_PASSWORD", ADMIN_TEST_PASSWORD)

    for module_name in list(sys.modules):
        if module_name == "app" or module_name.startswith("app."):
            del sys.modules[module_name]

    main = importlib.import_module("app.main")
    return TestClient(main.app), main


def test_injection_schedule_calculations_are_decimal_safe():
    assert parse_machine_a_value("2A半") == Decimal("2.5")
    assert parse_machine_a_value("50A 400T") == Decimal("50")
    assert parse_machine_a_value("无机安") is None
    assert normalize_ratio("8%") == Decimal("0.08")
    assert normalize_ratio("0.08") == Decimal("0.08")
    assert qualified_shots("100", "3") == Decimal("97")
    assert qualified_shots("2", "3") == Decimal(0)
    assert progress("0", "12") == (Decimal("12"), Decimal(0), Decimal(0))
    assert progress("200", "50") == (
        Decimal("50"),
        Decimal("150"),
        Decimal("25.00"),
    )
    assert estimated_material_kg("1000", "24", "8%") == Decimal("25.920")
    assert effective_changeover_hours("2.5", None) == Decimal("2.5")
    assert effective_changeover_hours("2.5", "1.25") == Decimal("1.25")


def test_read_only_api_is_factory_scoped_and_returns_backend_progress(monkeypatch):
    client, main = make_client(monkeypatch)
    with client:
        login = client.post(
            "/api/auth/login",
            json={"username": "admin", "password": ADMIN_TEST_PASSWORD},
        )
        assert login.status_code == 200

        models = importlib.import_module("app.models.injection_schedule")
        db_module = importlib.import_module("app.db")
        now = "2026-08-30T08:00:00+08:00"
        with db_module.SessionLocal() as db:
            for factory_id, suffix in (("huaxing", "hx"), ("huakang-b", "hkb")):
                order = models.InjectionScheduleOrderDemand(
                    id=f"order-{suffix}",
                    factory_id=factory_id,
                    order_no=f"ORDER-{suffix}",
                    product_code=f"ITEM-{suffix}",
                    mold_code=f"MOLD-{suffix}",
                    quantity_sets=Decimal("100"),
                    order_shots=Decimal("100"),
                    delivery_due_date="2026-09-10",
                    data_completeness_status="COMPLETE",
                    created_by="test",
                    updated_by="test",
                    created_at=now,
                    updated_at=now,
                )
                machine = models.InjectionScheduleMachine(
                    id=f"machine-{suffix}",
                    factory_id=factory_id,
                    machine_code=f"M-{suffix}",
                    machine_a_label="24A 260T",
                    machine_ounce_capacity=Decimal("24"),
                    status="AVAILABLE",
                    created_by="test",
                    updated_by="test",
                    created_at=now,
                    updated_at=now,
                )
                db.add_all((order, machine))
                db.flush()
                line = models.InjectionScheduleLine(
                    id=f"line-{suffix}",
                    factory_id=factory_id,
                    order_demand_id=order.id,
                    machine_id=machine.id,
                    sequence_no=Decimal(0),
                    status="SCHEDULED",
                    planned_start_at="2026-08-30T08:00:00+08:00",
                    planned_finish_at="2026-08-31T08:00:00+08:00",
                    created_by="test",
                    updated_by="test",
                    created_at=now,
                    updated_at=now,
                )
                db.add(line)
                db.flush()
            db.add(
                models.InjectionScheduleShiftOutput(
                    id="output-hx",
                    factory_id="huaxing",
                    schedule_line_id="line-hx",
                    production_date="2026-08-30",
                    shift="DAY",
                    reported_shots=Decimal("30"),
                    defect_shots=Decimal("2"),
                    qualified_shots=Decimal("28"),
                    reported_by="test",
                    reported_at=now,
                    updated_at=now,
                )
            )
            db.commit()

        bootstrap = client.get(
            "/api/production/injection-scheduling/bootstrap",
            params={"factory_id": "huaxing"},
        )
        assert bootstrap.status_code == 200
        assert bootstrap.json()["capabilities"] == {
            "can_read": True,
            "can_edit": True,
            "can_schedule": True,
            "can_admin": True,
            "phase": "LOCAL_ACCEPTANCE",
        }

        orders = client.get(
            "/api/production/injection-scheduling/orders",
            params={"factory_id": "huaxing"},
        )
        assert orders.status_code == 200
        assert orders.json()["total"] == 1
        assert orders.json()["items"][0]["order_no"] == "ORDER-hx"
        assert orders.json()["items"][0]["qualified_shots"] == "28.000000"
        assert orders.json()["items"][0]["remaining_shots"] == "72.000000"

        other_factory = client.get(
            "/api/production/injection-scheduling/orders",
            params={"factory_id": "huakang-b"},
        )
        assert other_factory.status_code == 200
        assert [item["order_no"] for item in other_factory.json()["items"]] == [
            "ORDER-hkb"
        ]

        board = client.get(
            "/api/production/injection-scheduling/board",
            params={
                "factory_id": "huaxing",
                "start_date": "2026-08-30",
                "end_date": "2026-08-31",
            },
        )
        assert board.status_code == 200
        assert board.json()["kpis"]["active_order_count"] == 1
        assert board.json()["lines"][0]["machine_code"] == "M-hx"

        unsupported = client.get(
            "/api/production/injection-scheduling/bootstrap",
            params={"factory_id": "group"},
        )
        assert unsupported.status_code == 422

    assert main.app is not None


def _flat_order_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "仓库订单"
    sheet.append(
        [
            "款号",
            "模具编号",
            "工模名称",
            "单号",
            "数量",
            "总套数",
            "啤数",
            "机安",
            "颜色",
            "色粉号",
            "用料名称",
            "水口比例",
            "模具日产量",
            "整啤净重",
            "交货日期",
            "备注",
        ]
    )
    sheet.append(
        [
            "000123",
            "MOLD-001",
            "测试车壳",
            "ORDER-001",
            1000,
            1000,
            500,
            "24A",
            "蓝色",
            "00428",
            "ABS 750NSW",
            "8%",
            2000,
            24,
            "2026-09-10",
            "喷油优先",
        ]
    )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _plan_history_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "计划表"
    sheet.cell(1, 1, "河源华兴啤机生产日计划表")
    sheet.cell(2, 16, datetime(2026, 8, 29))
    headers = [
        "机位",
        "工模",
        "名称",
        "单号",
        "货号",
        "数量",
        "啤数",
        "机安",
        "交货日期",
        "模具日产量",
        "水口比例",
        "用料",
        "颜色",
        "备注",
        "占位",
        "白班",
        "夜班",
    ]
    for column, value in enumerate(headers, start=1):
        sheet.cell(3, column, value)
    sheet.append(["1号", None, None, None, None, None, None, "24A"])
    sheet.append(
        [
            None,
            "MOLD-PLAN-1",
            "历史订单",
            "ORDER-PLAN-1",
            "000789",
            500,
            500,
            "24A",
            datetime(2026, 9, 10),
            1000,
            "8%",
            "ABS",
            "蓝色",
            "",
            "",
            120,
            80,
        ]
    )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_excel_preview_is_deterministic_and_commit_is_idempotent(monkeypatch):
    client, _ = make_client(monkeypatch)
    content = _flat_order_workbook()
    with client:
        login = client.post(
            "/api/auth/login",
            json={"username": "admin", "password": ADMIN_TEST_PASSWORD},
        )
        assert login.status_code == 200

        preview = client.post(
            "/api/production/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "request_id": "preview-flat-1"},
            files={
                "file": (
                    "warehouse-orders.xlsx",
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert preview.status_code == 200, preview.text
        payload = preview.json()
        assert payload["profile_code"] == "WAREHOUSE_ORDER_FLAT_V1"
        assert payload["status"] == "READY"
        assert payload["summary"] == {
            "row_count": 1,
            "machine_count": 0,
            "history_output_count": 0,
            "blocking_issue_count": 0,
            "warning_count": 0,
            "create_count": 1,
            "update_count": 0,
            "unchanged_count": 0,
        }
        assert payload["rows"][0]["change_action"] == "CREATE"
        assert payload["rows"][0]["product_code"] == "000123"
        assert payload["rows"][0]["pigment_code"] == "00428"
        assert payload["rows"][0]["order_shots"] == "500"
        assert payload["rows"][0]["water_ratio"] == "0.08"

        duplicate_preview = client.post(
            "/api/production/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "request_id": "preview-flat-1"},
            files={"file": ("warehouse-orders.xlsx", content)},
        )
        assert duplicate_preview.status_code == 200
        assert duplicate_preview.json()["batch_id"] == payload["batch_id"]

        commit = client.post(
            f"/api/production/injection-scheduling/imports/{payload['batch_id']}/commit",
            json={"factory_id": "huaxing", "request_id": "commit-flat-1"},
        )
        assert commit.status_code == 200, commit.text
        assert commit.json()["created_count"] == 1

        repeated_commit = client.post(
            f"/api/production/injection-scheduling/imports/{payload['batch_id']}/commit",
            json={"factory_id": "huaxing", "request_id": "commit-flat-1"},
        )
        assert repeated_commit.status_code == 200
        assert repeated_commit.json() == commit.json()

        orders = client.get(
            "/api/production/injection-scheduling/orders",
            params={"factory_id": "huaxing"},
        )
        assert orders.status_code == 200
        assert orders.json()["total"] == 1
        assert orders.json()["items"][0]["product_code"] == "000123"


def test_one_unified_template_serves_four_factories_and_enforces_factory_scope(
    monkeypatch,
):
    client, _ = make_client(monkeypatch)
    factory_ids = ("huakang-a", "huakang-b", "huadeng", "huaxing")
    expected_content: bytes | None = None
    huaxing_content = b""
    with client:
        assert (
            client.post(
                "/api/auth/login",
                json={"username": "admin", "password": ADMIN_TEST_PASSWORD},
            ).status_code
            == 200
        )
        for factory_id in factory_ids:
            response = client.get(
                "/api/production/injection-scheduling/templates/unified-plan",
                params={"factory_id": factory_id},
            )
            assert response.status_code == 200, response.text
            assert response.content[:2] == b"PK"
            if expected_content is None:
                expected_content = response.content
            assert response.content == expected_content
            workbook = load_workbook(BytesIO(response.content), read_only=True)
            try:
                assert workbook.sheetnames == ["统一计划表"]
                plan = workbook["统一计划表"]
                assert plan["A1"].value == "Royal Regent Nexus · 统一注塑排产计划表"
                assert plan["A2"].value == "模板版本：2.0"
                assert (
                    tuple(plan.cell(4, column).value for column in range(1, 57))
                    == UNIFIED_PLAN_HEADERS
                )
            finally:
                workbook.close()
            if factory_id == "huaxing":
                huaxing_content = response.content

        for unsupported_factory_id in ("huakang-c", "huakang-d"):
            response = client.get(
                "/api/production/injection-scheduling/templates/unified-plan",
                params={"factory_id": unsupported_factory_id},
            )
            assert response.status_code == 422
            assert "没有注塑部" in response.text

        workbook = load_workbook(BytesIO(huaxing_content))
        plan = workbook["统一计划表"]
        values = {
            3: "待排",
            5: "特急",
            6: "24A",
            7: "MOLD-UNIFIED-001",
            8: "统一模板订单",
            9: "ORDER-UNIFIED-001",
            10: "000456",
            11: 1000,
            12: 500,
            16: 2000,
            18: 0.08,
            19: "深蓝色",
            20: "00428",
            21: "深",
            22: "ABS 750NSW",
            23: 24,
            28: datetime(2026, 9, 15),
            46: "已齐料",
            47: 10,
            51: "huaxing",
        }
        for column, value in values.items():
            plan.cell(5, column, value)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        filled_content = output.getvalue()

        preview = client.post(
            "/api/production/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "request_id": "preview-unified-1"},
            files={"file": ("huaxing-unified.xlsx", filled_content)},
        )
        assert preview.status_code == 200, preview.text
        payload = preview.json()
        assert payload["profile_code"] == "UNIFIED_PLAN_V2"
        assert payload["source_sheet"] == "统一计划表"
        assert payload["summary"]["row_count"] == 1
        assert payload["summary"]["blocking_issue_count"] == 0
        assert payload["rows"][0]["priority"] == "EXPEDITE"
        assert payload["rows"][0]["product_code"] == "000456"

        mismatch = client.post(
            "/api/production/injection-scheduling/imports/preview",
            data={"factory_id": "huakang-a", "request_id": "preview-unified-2"},
            files={"file": ("huaxing-unified.xlsx", filled_content)},
        )
        assert mismatch.status_code == 200
        mismatch_payload = mismatch.json()
        assert mismatch_payload["summary"]["blocking_issue_count"] == 1
        assert mismatch_payload["issues"][0]["error_type"] == "FACTORY_SCOPE_MISMATCH"

        workbook = load_workbook(BytesIO(huaxing_content))
        workbook.create_sheet("机台主数据")
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        invalid_multi_sheet = client.post(
            "/api/production/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "request_id": "preview-unified-3"},
            files={"file": ("invalid-multi-sheet.xlsx", output.getvalue())},
        )
        assert invalid_multi_sheet.status_code == 422
        assert "只能包含" in invalid_multi_sheet.text


def test_molds_are_company_shared_while_machines_remain_factory_scoped(monkeypatch):
    client, _ = make_client(monkeypatch)
    machine_payload = {
        "machine_code": "M-01",
        "position": "A区",
        "machine_name": "一号机",
        "machine_a_label": "24A",
        "machine_ounce_capacity": 24,
    }
    mold_payload = {
        "factory_id": "huaxing",
        "mold_code": "MOLD-SHARED-001",
        "product_code": "000789",
        "product_name": "共享模具产品",
        "required_machine_a_label": "24A",
        "mold_ounce_requirement": 24,
    }
    with client:
        assert (
            client.post(
                "/api/auth/login",
                json={"username": "admin", "password": ADMIN_TEST_PASSWORD},
            ).status_code
            == 200
        )
        for factory_id in ("huaxing", "huadeng"):
            response = client.post(
                "/api/production/injection-scheduling/machines",
                json={"factory_id": factory_id, **machine_payload},
            )
            assert response.status_code == 200, response.text

        created_mold = client.post(
            "/api/production/injection-scheduling/molds", json=mold_payload
        )
        assert created_mold.status_code == 200, created_mold.text
        assert created_mold.json()["factory_id"] == "company"
        assert created_mold.json()["scope_type"] == "COMPANY_SHARED"

        for factory_id in ("huaxing", "huadeng", "huakang-a", "huakang-b"):
            molds = client.get(
                "/api/production/injection-scheduling/molds",
                params={"factory_id": factory_id},
            )
            assert molds.status_code == 200, molds.text
            assert molds.json()["total"] == 1
            assert molds.json()["items"][0]["id"] == created_mold.json()["id"]
            assert molds.json()["items"][0]["scope_type"] == "COMPANY_SHARED"

        duplicate = client.post(
            "/api/production/injection-scheduling/molds",
            json={**mold_payload, "factory_id": "huadeng"},
        )
        assert duplicate.status_code == 409

        huaxing_machines = client.get(
            "/api/production/injection-scheduling/machines",
            params={"factory_id": "huaxing"},
        ).json()["items"]
        huadeng_machines = client.get(
            "/api/production/injection-scheduling/machines",
            params={"factory_id": "huadeng"},
        ).json()["items"]
        assert len(huaxing_machines) == len(huadeng_machines) == 1
        assert huaxing_machines[0]["id"] != huadeng_machines[0]["id"]


def test_plan_history_preview_and_commit_preserve_day_night_outputs(monkeypatch):
    client, _ = make_client(monkeypatch)
    content = _plan_history_workbook()
    with client:
        assert (
            client.post(
                "/api/auth/login",
                json={"username": "admin", "password": ADMIN_TEST_PASSWORD},
            ).status_code
            == 200
        )
        preview = client.post(
            "/api/production/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "request_id": "preview-history-1"},
            files={"file": ("plan-history.xlsx", content)},
        )
        assert preview.status_code == 200, preview.text
        payload = preview.json()
        assert payload["profile_code"] == "HUA_XING_PLAN_V1"
        assert payload["summary"]["history_output_count"] == 2
        assert payload["summary"]["blocking_issue_count"] == 0
        assert payload["history_outputs"] == [
            {
                "business_key": payload["rows"][0]["business_key"],
                "source_row": 5,
                "production_date": "2026-08-29",
                "shift": "DAY",
                "reported_shots": "120",
            },
            {
                "business_key": payload["rows"][0]["business_key"],
                "source_row": 5,
                "production_date": "2026-08-29",
                "shift": "NIGHT",
                "reported_shots": "80",
            },
        ]

        commit = client.post(
            f"/api/production/injection-scheduling/imports/{payload['batch_id']}/commit",
            json={"factory_id": "huaxing", "request_id": "commit-history-1"},
        )
        assert commit.status_code == 200, commit.text
        assert commit.json()["history_created_count"] == 2

        db_module = importlib.import_module("app.db")
        models = importlib.import_module("app.models.injection_schedule")
        with db_module.SessionLocal() as db:
            outputs = db.scalars(
                select(models.InjectionScheduleShiftOutput).order_by(
                    models.InjectionScheduleShiftOutput.shift
                )
            ).all()
            assert [(item.shift, item.qualified_shots) for item in outputs] == [
                ("DAY", Decimal("120")),
                ("NIGHT", Decimal("80")),
            ]


def test_manual_move_shift_output_and_optimistic_conflict(monkeypatch):
    client, main = make_client(monkeypatch)
    content = _flat_order_workbook()
    with client:
        assert (
            client.post(
                "/api/auth/login",
                json={"username": "admin", "password": ADMIN_TEST_PASSWORD},
            ).status_code
            == 200
        )
        preview = client.post(
            "/api/production/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "request_id": "preview-manual-1"},
            files={"file": ("warehouse-orders.xlsx", content)},
        ).json()
        commit = client.post(
            f"/api/production/injection-scheduling/imports/{preview['batch_id']}/commit",
            json={"factory_id": "huaxing", "request_id": "commit-manual-1"},
        )
        assert commit.status_code == 200, commit.text

        from app.db import SessionLocal
        from app.models.injection_schedule import InjectionScheduleMachine

        with SessionLocal() as db:
            now = "2026-08-30T12:00:00+08:00"
            db.add(
                InjectionScheduleMachine(
                    id="machine-manual-1",
                    factory_id="huaxing",
                    machine_code="HX-01",
                    position="1",
                    machine_name="测试机",
                    machine_a_label="24A",
                    machine_ounce_capacity=Decimal("24"),
                    status="AVAILABLE",
                    created_by="test",
                    updated_by="test",
                    created_at=now,
                    updated_at=now,
                )
            )
            db.commit()

        board = client.get(
            "/api/production/injection-scheduling/board",
            params={
                "factory_id": "huaxing",
                "start_date": "2026-08-30",
                "end_date": "2026-09-15",
            },
        ).json()
        line = board["lines"][0]
        move_payload = {
            "factory_id": "huaxing",
            "target_machine_id": "machine-manual-1",
            "planned_start_at": "2026-09-01T08:00:00+08:00",
            "planned_finish_at": "2026-09-02T08:00:00+08:00",
            "expected_line_version": line["version"],
            "expected_schedule_revision": board["schedule_revision"],
            "reason": "测试人工排产",
        }
        validation = client.post(
            f"/api/production/injection-scheduling/schedule-lines/{line['id']}/validate-move",
            json=move_payload,
        )
        assert validation.status_code == 200, validation.text
        assert validation.json()["valid"] is True
        move_payload["validate_token"] = validation.json()["validate_token"]
        moved = client.post(
            f"/api/production/injection-scheduling/schedule-lines/{line['id']}/move",
            json=move_payload,
        )
        assert moved.status_code == 200, moved.text
        assert moved.json()["schedule_revision"] == 2

        output = client.put(
            f"/api/production/injection-scheduling/schedule-lines/{line['id']}/shift-outputs/2026-09-01/DAY",
            json={
                "factory_id": "huaxing",
                "reported_shots": "120",
                "defect_shots": "5",
                "downtime_minutes": 15,
            },
        )
        assert output.status_code == 200, output.text
        assert output.json()["qualified_shots"] == 115.0
        assert output.json()["schedule_revision"] == 3

        stale = client.post(
            f"/api/production/injection-scheduling/schedule-lines/{line['id']}/validate-move",
            json=move_payload,
        )
        assert stale.status_code == 409


def test_auto_schedule_prefers_exact_machine_and_rejects_stale_digest(monkeypatch):
    client, _ = make_client(monkeypatch)
    with client:
        assert (
            client.post(
                "/api/auth/login",
                json={"username": "admin", "password": ADMIN_TEST_PASSWORD},
            ).status_code
            == 200
        )
        content = _flat_order_workbook()
        preview = client.post(
            "/api/production/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "request_id": "preview-auto-1"},
            files={"file": ("warehouse-orders.xlsx", content)},
        ).json()
        assert (
            client.post(
                f"/api/production/injection-scheduling/imports/{preview['batch_id']}/commit",
                json={"factory_id": "huaxing", "request_id": "commit-auto-1"},
            ).status_code
            == 200
        )

        from app.db import SessionLocal
        from app.models.injection_schedule import (
            InjectionScheduleMachine,
            InjectionScheduleOrderDemand,
        )

        with SessionLocal() as db:
            order = (
                db.query(InjectionScheduleOrderDemand)
                .filter_by(factory_id="huaxing")
                .one()
            )
            order.material_status = "READY"
            now = "2026-08-30T12:00:00+08:00"
            for machine_id, machine_code, capacity in (
                ("machine-exact", "HX-24A", Decimal("24")),
                ("machine-large", "HX-40A", Decimal("40")),
            ):
                db.add(
                    InjectionScheduleMachine(
                        id=machine_id,
                        factory_id="huaxing",
                        machine_code=machine_code,
                        machine_a_label=f"{capacity}A",
                        machine_ounce_capacity=capacity,
                        status="AVAILABLE",
                        created_by="test",
                        updated_by="test",
                        created_at=now,
                        updated_at=now,
                    )
                )
            db.commit()

        proposal = client.post(
            "/api/production/injection-scheduling/auto-schedule/preview",
            json={
                "factory_id": "huaxing",
                "request_id": "auto-preview-1",
                "start_at": "2026-09-01T08:00:00+08:00",
                "end_at": "2026-09-15T08:00:00+08:00",
                "mode": "INCREMENTAL",
                "selected_order_ids": [],
                "affected_machine_ids": [],
                "expected_schedule_revision": 1,
            },
        )
        assert proposal.status_code == 200, proposal.text
        payload = proposal.json()
        assert payload["summary"]["scheduled_count"] == 1
        assert payload["changes"][0]["machine_id"] == "machine-exact"

        with SessionLocal() as db:
            order = (
                db.query(InjectionScheduleOrderDemand)
                .filter_by(factory_id="huaxing")
                .one()
            )
            order.version += 1
            db.commit()

        stale = client.post(
            f"/api/production/injection-scheduling/auto-schedule/proposals/{payload['proposal_id']}/apply",
            json={
                "factory_id": "huaxing",
                "expected_schedule_revision": 1,
                "reason": "测试失效保护",
            },
        )
        assert stale.status_code == 409
        assert "重新生成预览" in stale.text

        fresh_request = {
            "factory_id": "huaxing",
            "request_id": "auto-preview-2",
            "start_at": "2026-09-01T08:00:00+08:00",
            "end_at": "2026-09-15T08:00:00+08:00",
            "mode": "INCREMENTAL",
            "selected_order_ids": [],
            "affected_machine_ids": [],
            "expected_schedule_revision": 1,
        }
        fresh = client.post(
            "/api/production/injection-scheduling/auto-schedule/preview",
            json=fresh_request,
        )
        assert fresh.status_code == 200, fresh.text
        applied = client.post(
            f"/api/production/injection-scheduling/auto-schedule/proposals/{fresh.json()['proposal_id']}/apply",
            json={
                "factory_id": "huaxing",
                "expected_schedule_revision": 1,
                "reason": "确认自动排程",
            },
        )
        assert applied.status_code == 200, applied.text
        assert applied.json()["status"] == "APPLIED"

        board = client.get(
            "/api/production/injection-scheduling/board",
            params={
                "factory_id": "huaxing",
                "start_date": "2026-09-01",
                "end_date": "2026-09-15",
            },
        ).json()
        scheduled_line = board["lines"][0]
        assert scheduled_line["machine_id"] == "machine-exact"
        locked = client.post(
            f"/api/production/injection-scheduling/schedule-lines/{scheduled_line['id']}/lock",
            json={
                "factory_id": "huaxing",
                "expected_line_version": scheduled_line["version"],
                "expected_schedule_revision": board["schedule_revision"],
                "reason": "冻结已确认任务",
            },
        )
        assert locked.status_code == 200, locked.text
        locked_preview = client.post(
            "/api/production/injection-scheduling/auto-schedule/preview",
            json={
                **fresh_request,
                "request_id": "auto-preview-locked",
                "mode": "FULL",
                "expected_schedule_revision": locked.json()["schedule_revision"],
            },
        )
        assert locked_preview.status_code == 200, locked_preview.text
        assert locked_preview.json()["summary"]["scheduled_count"] == 0
        assert locked_preview.json()["unscheduled"][0]["reason_code"] == "LOCKED"


def test_saved_view_export_and_disabled_ai_do_not_block_core(monkeypatch):
    monkeypatch.setenv("QWEN_ENABLED", "false")
    client, _ = make_client(monkeypatch)
    with client:
        assert (
            client.post(
                "/api/auth/login",
                json={"username": "admin", "password": ADMIN_TEST_PASSWORD},
            ).status_code
            == 200
        )
        preview = client.post(
            "/api/production/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "request_id": "preview-export-1"},
            files={"file": ("warehouse-orders.xlsx", _flat_order_workbook())},
        ).json()
        assert (
            client.post(
                f"/api/production/injection-scheduling/imports/{preview['batch_id']}/commit",
                json={"factory_id": "huaxing", "request_id": "commit-export-1"},
            ).status_code
            == 200
        )

        saved = client.post(
            "/api/production/injection-scheduling/saved-views",
            json={
                "factory_id": "huaxing",
                "name": "急单视图",
                "scope": "PERSONAL",
                "config": {
                    "columns": ["order_no", "product_code"],
                    "filters": {"priority": "URGENT"},
                },
            },
        )
        assert saved.status_code == 200, saved.text
        listed = client.get(
            "/api/production/injection-scheduling/saved-views",
            params={"factory_id": "huaxing"},
        )
        assert listed.status_code == 200
        assert listed.json()[0]["name"] == "急单视图"

        exported = client.post(
            "/api/production/injection-scheduling/exports/table",
            json={
                "factory_id": "huaxing",
                "columns": ["order_no", "product_code", "remaining_shots"],
                "search": "ORDER-001",
            },
        )
        assert exported.status_code == 200, exported.text
        assert exported.content.startswith(b"PK")
        assert "spreadsheetml" in exported.headers["content-type"]

        ai = client.post(
            "/api/production/injection-scheduling/ai/interpret-filter",
            json={"factory_id": "huaxing", "text": "查看所有的急单"},
        )
        assert ai.status_code == 503

        orders = client.get(
            "/api/production/injection-scheduling/orders",
            params={"factory_id": "huaxing"},
        )
        assert orders.status_code == 200
        assert orders.json()["total"] == 1
