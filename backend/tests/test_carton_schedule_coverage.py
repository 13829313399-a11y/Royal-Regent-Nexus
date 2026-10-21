"""Scoped absence decisions, complete input evidence and import-task replay."""
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier, Event
from io import BytesIO
from uuid import uuid4
import pickle
import time

import pytest
from openpyxl import load_workbook

from test_molding_sample_api import make_client, login_as
from test_carton_item_schedule import workbook, source_row, _prepare_schedule_customer

BASE = "/api/carton-procurement"
COVERAGE = BASE + "/schedule-coverage"


@pytest.fixture(autouse=True)
def isolated(monkeypatch, tmp_path):
    import test_molding_sample_api
    monkeypatch.setattr(test_molding_sample_api, "TEST_TMP_DIR", tmp_path)


def fields(mode="FULL", day="2026-10-10"):
    result = dict(customer_code="DICKIE", scope_mode=mode, as_of_date=day)
    if mode == "FULL": result.update(coverage_start="2026-09-01", coverage_end="2026-09-30", full_scope_confirmed="true")
    return result


def upload(client, content, *, mode="FULL", day="2026-10-10", expected=201):
    response = client.post(BASE + "/weekly-imports", params={"factory_id": "huaxing"},
        data=fields(mode, day), files={"file": ("scope.xlsx", content)})
    assert response.status_code == expected, response.text
    return response.json()


def workspace(client, status="ALL", **params):
    response = client.get(COVERAGE + "/workspace", params=dict(factory_id="huaxing", status=status, **params))
    assert response.status_code == 200, response.text
    return response.json()


def body(row, decision="KEEP_REQUIRED"):
    return dict(factory_id="huaxing", request_id=uuid4().hex, expected_revision=row["revision"],
        source_fingerprint=row["source_fingerprint"], decision=decision, reason="核实客户完整排期及原采购需求")


def test_unknown_sheets_are_visible_in_both_structured_diagnostic_and_warnings():
    from app.services.carton_procurement_imports import _parse_weekly
    book = load_workbook(BytesIO(workbook([source_row(contract="KNOWN")])))
    book.create_sheet("额外业务明细").append(["OTHER", 10])
    output = BytesIO(); book.save(output)
    parsed = _parse_weekly("diagnostic.xlsx", output.getvalue())
    assert parsed["coverage_parse"]["unrecognized_sheets"] == ["额外业务明细"]
    assert len(parsed["rows"]) == 1 and parsed["rows"][0]["contract_no"] == "KNOWN"
    assert any("额外业务明细" in warning and "不是全量覆盖" in warning for warning in parsed["warnings"])
    clean = _parse_weekly("clean.xlsx", workbook([source_row(contract="KNOWN")]))
    assert not clean["coverage_parse"]["unrecognized_sheets"]
    assert not any("不是全量覆盖" in warning for warning in clean["warnings"])


def test_concurrent_review_has_one_winner_old_tasks_supersede_and_undo_restores(monkeypatch):
    with make_client(monkeypatch) as client:
        _prepare_schedule_customer(client)
        baseline = upload(client, workbook([source_row(contract="MISSING"), source_row(contract="STAYS")]))
        first = upload(client, workbook([source_row(contract="STAYS")]), day="2026-10-11")
        row = workspace(client)["missing"]["items"][0]
        barrier = Barrier(2)
        commands = [body(row, name) for name in ("KEEP_REQUIRED", "SOURCE_OMISSION")]
        def submit(index):
            barrier.wait(timeout=10)
            return client.post(COVERAGE + f"/missing/{row['id']}/review", json=commands[index])
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(submit, (0, 1)))
        assert sorted(result.status_code for result in responses) == [200, 409]
        winner = next(index for index, result in enumerate(responses) if result.status_code == 200)
        replayed = client.post(COVERAGE + f"/missing/{row['id']}/review", json=commands[winner])
        assert replayed.status_code == 200 and replayed.json()["replayed"]
        assert replayed.json()["revision"] == row["revision"] + 1
        newer = upload(client, workbook([source_row(contract="STAYS", qty=400)]), day="2026-10-12")
        pending = workspace(client, "PENDING")["missing"]["items"]
        assert len(pending) == 1 and pending[0]["id"] != row["id"]
        old = client.get(COVERAGE + f"/missing/{row['id']}", params={"factory_id": "huaxing"}).json()
        assert old["status"] == "SUPERSEDED" and not old["can_review"]
        assert client.post(COVERAGE + f"/missing/{row['id']}/review", json=body(old)).status_code == 409
        appeared = upload(client, workbook([source_row("备料单", contract="MISSING"), source_row(contract="STAYS")]),
            mode="INCREMENTAL", day="2026-10-13")
        returned = client.get(COVERAGE + f"/missing/{pending[0]['id']}", params={"factory_id": "huaxing"}).json()
        assert returned["status"] == "RETURNED" and not returned["can_operate_source"]
        old_source = next(item for item in baseline["parse_summary"]["rows"] if item["contract_no"] == "MISSING")
        forbidden = client.post(BASE + "/schedule-order-marks", json=dict(factory_id="huaxing", batch_id=baseline["id"],
            source_sheet=old_source["source_sheet"], source_row=old_source["source_row"], marked=True))
        assert forbidden.status_code == 409, forbidden.text
        assert client.post(BASE + f"/imports/{appeared['id']}/undo", json=dict(factory_id="huaxing", reason="撤销业务类型填错的来源")).status_code == 200
        restored = client.get(COVERAGE + f"/missing/{pending[0]['id']}", params={"factory_id": "huaxing"}).json()
        assert restored["status"] == "PENDING" and not restored["can_operate_source"]
        undone = client.post(BASE + f"/imports/{newer['id']}/undo", json=dict(factory_id="huaxing", reason="撤销误上传的全量排期"))
        assert undone.status_code == 200, undone.text
        old = client.get(COVERAGE + f"/missing/{row['id']}", params={"factory_id": "huaxing"}).json()
        assert old["status"] in {"KEEP_REQUIRED", "SOURCE_OMISSION"} and old["can_review"]
        current = client.get(BASE + f"/imports/{baseline['id']}", params={"factory_id": "huaxing"}).json()
        source = next(item for item in current["parse_summary"]["rows"] if item["contract_no"] == "MISSING")
        assert source["can_operate_source"] and source["missing_review"]["id"] == row["id"]
        summary = client.get(BASE + "/imports", params=dict(factory_id="huaxing", import_type="WEEKLY_SCHEDULE", summary_only=True)).json()
        compact = next(item for item in summary["items"] if item["id"] == first["id"])["parse_summary"]["schedule_coverage"]
        assert compact["missing_count"] == 1 and "missing_rows" not in compact and "identity_profile" not in compact


def test_presence_all_rows_duplicate_snapshots_and_unknown_dates_are_conservative(monkeypatch):
    with make_client(monkeypatch) as client:
        _prepare_schedule_customer(client)
        unknown = source_row(contract="UNKNOWN"); unknown[-1] = "待核实"
        upload(client, workbook([source_row(contract="DUP"), source_row(contract="DUP", qty=200),
            source_row(contract="NONFORMAL"), unknown, source_row(contract="STAYS")]))
        current = workbook([source_row("备料单", contract="NONFORMAL", qty=0), source_row(contract="STAYS")])
        imported = upload(client, current, day="2026-10-11")
        rows = workspace(client, "PENDING")["missing"]["items"]
        assert [row["contract_no"] for row in rows] == ["DUP"]
        assert len(rows[0]["source_rows"]) == 2
        assert {row["quantity"] for row in rows[0]["source_rows"]} == {200, 360}
        assert any(row["contract_no"] == "UNKNOWN" for row in imported["parse_summary"]["schedule_coverage"]["excluded_rows"])
        book = load_workbook(BytesIO(current)); book.create_sheet("未识别来源").append(["BROKEN", 10])
        output = BytesIO(); book.save(output)
        upload(client, output.getvalue(), day="2026-10-12", expected=422)
        incremental = upload(client, output.getvalue(), mode="INCREMENTAL", day="2026-10-12")
        assert incremental["parse_summary"]["coverage_parse"]["unrecognized_sheets"] == ["未识别来源"]
        assert len(workspace(client, limit=1)["customers"]) == len(workspace(client, limit=100)["customers"])


def test_file_job_scope_identity_complete_replay_and_current_authorization(monkeypatch):
    with make_client(monkeypatch) as client:
        _prepare_schedule_customer(client)
        from app.services import carton_file_jobs as jobs
        from app.services.carton_procurement_imports import parse_carton_file
        def compute(folder, operation, args):
            parsed = parse_carton_file(*args)
            with (folder / "input.pkl").open("wb") as stream: pickle.dump((operation, args), stream)
            with (folder / "result.pkl").open("wb") as stream: pickle.dump({"value": parsed}, stream)
            return parsed
        monkeypatch.setattr(jobs, "_compute", compute)
        content = workbook([source_row(contract="ASYNC")])
        def start(mode, day):
            result = client.post(BASE + "/file-jobs/imports", params=dict(factory_id="huaxing", import_type="WEEKLY_SCHEDULE"),
                data=fields(mode, day), files={"file": ("async.xlsx", content)})
            assert result.status_code == 202, result.text
            identifier = result.json()["id"]
            deadline = time.monotonic() + 5
            while time.monotonic() < deadline:
                if client.get(BASE + f"/file-jobs/{identifier}", params={"factory_id": "huaxing"}).json()["status"] == "READY": return identifier
                Event().wait(.01)
            raise AssertionError("synthetic parser did not finish")
        try:
            first = start("FULL", "2026-10-10")
            url = BASE + f"/file-jobs/{first}/complete"
            done = client.post(url, params={"factory_id": "huaxing"})
            assert done.status_code == 200, done.text
            row = done.json()
            assert len(row["import_profile"]) < 255 and row["import_profile"].startswith("schedule-coverage-v1:")
            assert row["parse_summary"]["schedule_coverage"]["scope_mode"] == "FULL"
            assert client.post(url, params={"factory_id": "huaxing"}).json()["id"] == row["id"]
            second = start("INCREMENTAL", "2026-10-10")
            changed = client.post(BASE + f"/file-jobs/{second}/complete", params={"factory_id": "huaxing"})
            assert changed.status_code == 200 and changed.json()["id"] != row["id"]
            from app.services import carton_procurement as core
            from app.models.auth import AuthUser
            from sqlalchemy import select
            original = core._lock_receipt_factory
            def disable_after_lock(db, factory):
                original(db, factory)
                actor = db.get(AuthUser, row["imported_by"])
                actor.status = "disabled"; db.flush()
            with monkeypatch.context() as mutation:
                mutation.setattr(core, "_lock_receipt_factory", disable_after_lock)
                refused = client.post(url, params={"factory_id": "huaxing"})
            assert refused.status_code == 403, refused.text
        finally:
            for identifier in list(jobs._jobs): jobs._expire(identifier)
