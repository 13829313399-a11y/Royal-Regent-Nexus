from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.services.internal_quote import _rr2_cost_summary
from app.services.internal_quote_calculator import calculate_section
from app.services.internal_quote_excel import _build_summary_sheet
from test_internal_quote_excel_formula_links import _section
from test_internal_quote_purchase_export import arithmetic_value


def export_costs(payloads, mode, components=None):
    snapshot = {'fx': {'rmb_hkd': '.85', 'hkd_usd': '7.8', 'rmb_usd': '6.63'}}
    sections, context = [], {}
    for code, payload in payloads.items():
        calc = calculate_section(code, payload, snapshot, 'COST-OWNERSHIP')
        assert calc['status'] == 'valid', calc.get('warnings')
        section = _section(code, payload=payload, calculation=calc)
        section.is_required = True
        sections.append(section)
        keys = ('hardware_hkd', 'auxiliary_hkd') if code == 'engineering' else ('total_hkd',)
        for key in keys:
            context[key if code == 'engineering' else f'{code}_hkd'] = Decimal(calc['totals'][key])
    context['factory_price_hkd'] = sum(context.values())
    sales = {'shipping': {'markup_x': '1.2', 'misc_ratio': '.02'}}
    if mode == 'component':
        sales.update(pricing_mode='component', pricing_components=components or [{'id': 'body', 'name': '公仔', 'markup_x': '1.15'}])
    sales_section = _section('sales', payload=sales)
    sales_section.is_required = True
    sections.append(sales_section)
    summary = _rr2_cost_summary(sections, context, snapshot,
                                factory_id='huakang-b' if mode == 'component' else 'huaxing')
    workbook = Workbook()
    _build_summary_sheet(workbook, SimpleNamespace(product_name='公仔', quote_no='COST-QA',
                         customer='JustPlay' if mode == 'component' else '普通客', region_code='mainland', remark=''),
                         sections, snapshot, summary, context, [])
    return workbook, context, summary


def row_amounts(sheet, category):
    return [arithmetic_value(sheet, f'D{row}') for row in range(1, sheet.max_row + 1)
            if sheet.cell(row, 2).value == category and sheet.cell(row, 3).value != '成本']


@pytest.mark.parametrize('mode', ['ordinary', 'detached', 'component'])
@pytest.mark.parametrize('legacy,standalone', [(False, True), (True, False), (True, True)])
def test_clothes_and_each_hair_source_export_once_without_negative_residual(mode, legacy, standalone):
    owner = {'pricing_component_id': 'body'}
    if mode == 'detached':
        owner['markup_override'] = '1.03'
    groups = [{'name': '测试公仔服装', 'category': 'clothes', **owner,
               'materials': [{'item': '布', 'usage': 1, 'unit_price_hkd': '5.9271',
                              'unit_price_source_currency': 'HKD', 'markup': 1}]}]
    if legacy:
        groups.append({'name': '旧车发', 'category': 'hair', **owner,
                       'materials': [{'item': '发丝', 'usage': 1, 'unit_price_rmb': '.85', 'markup': 1}]})
    payloads = {'sewing': {'groups': groups}}
    if standalone:
        payloads['hair'] = {'lines': [{'name': '车发', 'craft': '植发', 'weight_g': 2,
                                     'unit_price_hkd': '2.7642', 'unit': 'PCS', **owner}]}
    workbook, context, summary = export_costs(payloads, mode)
    try:
        sheet = workbook['报价明细']
        clothes = row_amounts(sheet, '车衣')
        hair = row_amounts(sheet, '车发')
        assert sum(clothes) == pytest.approx(5.9271)
        assert sum(hair) == pytest.approx((1 if legacy else 0) + (2.7642 if standalone else 0))
        assert all(value >= 0 for value in clothes + hair)
        assert sum(clothes + hair) == pytest.approx(float(context['factory_price_hkd']))
        assert not any(cell.value == '车缝未分类成本' for row in sheet for cell in row)
        assert float(next(row['value'] for row in summary['t1'] if row['key'] == 'sewing_hair')) == pytest.approx(sum(hair))
    finally:
        workbook.close()


@pytest.mark.parametrize('mode', ['ordinary', 'detached', 'component'])
def test_blister_names_do_not_duplicate_hardware_or_electronic_costs(mode):
    owner = {'pricing_component_id': 'body'}
    if mode == 'detached':
        owner['markup_override'] = '1.03'
    payloads = {
        'engineering': {'materials': [
            {'item': 'Blister 五金支架', 'category': 'hardware', 'quantity': 1, 'unit_price_rmb': '.85', **owner},
            {'item': '吸塑马达 motor', 'category': 'hardware', 'quantity': 1, 'unit_price_rmb': '.85', **owner},
            {'item': '真正吸塑', 'category': 'auxiliary', 'auxiliary_category': '吸塑',
             'quantity': 1, 'unit_price_rmb': '1.7', **owner},
        ]},
        'electronic': {'components': [{'item': 'Blister 电子组件', 'quantity': 1, 'unit_price_rmb': '.85', **owner}]},
    }
    workbook, context, summary = export_costs(payloads, mode)
    try:
        sheet = workbook['报价明细']
        values = {row['key']: Decimal(row['value']) for row in summary['t1']}
        assert values['suction'] == 2
        assert values['hardware'] == 1
        assert values['motor'] == 1
        assert values['electronic'] == context['electronic_hkd']
        assert sum(row_amounts(sheet, '吸塑')) == pytest.approx(2)
        # Component export retains its existing hardware/motor block grouping.
        amounts = [value for category in ('五金', '马达', '吸塑', '电子') for value in row_amounts(sheet, category)]
        assert all(value >= 0 for value in amounts)
        assert sum(amounts) == pytest.approx(float(context['factory_price_hkd']))
    finally:
        workbook.close()


@pytest.mark.parametrize('mode', ['ordinary', 'detached', 'component'])
def test_many_sewing_rows_preserve_authoritative_clothes_and_hair_subtotals(mode):
    owner = {'pricing_component_id': 'body'}
    if mode == 'detached':
        owner['markup_override'] = '1.03'
    payloads = {'sewing': {'groups': [
        {'name': f'{category}-{component}', 'category': category,
         'pricing_component_id': component,
         **({'markup_override': markup} if mode == 'detached' else {}),
         'materials': [{'item': f'物料{i}', 'usage': 1, 'unit_price_rmb': '.22', 'markup': 1}
                       for i in range(count)]}
        for category in ('clothes', 'hair')
        for component, count, markup in [('body', 12, '1.03'), ('accessory', 13, '1.07')]
    ]}, 'hair': {'lines': [{'name': '独立车发', 'craft': '植发', 'weight_g': 2,
                            'unit_price_hkd': '2.7642', 'unit': 'PCS', **owner}]}}
    workbook, context, summary = export_costs(payloads, mode, components=[
        {'id': 'body', 'name': '公仔', 'markup_x': '1.15'},
        {'id': 'accessory', 'name': '配件', 'markup_x': '1.25'},
    ])
    try:
        sheet = workbook['报价明细']
        expected = Decimal('6.4706')  # round(25 * .22 / .85, 4), not 25 rounded rows
        values = {row['key']: Decimal(row['value']) for row in summary['t1']}
        assert values['sewing_cloth'] == expected
        assert sum(row_amounts(sheet, '车衣')) == pytest.approx(float(expected))
        assert sum(row_amounts(sheet, '车发')) == pytest.approx(float(expected + Decimal('2.7642')))
        entries = summary['shipping_pricing']['pricing_entries']
        for category in ('clothes', 'hair'):
            assert sum(Decimal(e['amount_hkd']) for e in entries
                       if e['section'] == 'sewing' and e['category'] == category) == expected
        groups = summary['shipping_pricing']['pricing_groups']
        assert sum(Decimal(g['cost_hkd']) for g in groups) == context['factory_price_hkd']
        if mode == 'component':
            for group in groups:
                assert Decimal(group['cost_hkd']) == sum(Decimal(e['amount_hkd']) for e in entries
                                                         if e['pricing_component_id'] == group['id'])
    finally:
        workbook.close()


@pytest.mark.parametrize('mode', ['ordinary', 'detached', 'component'])
def test_two_categories_rounding_up_share_one_department_budget(mode):
    payloads = {'sewing': {'groups': [
        {'name': category, 'category': category, 'pricing_component_id': 'body',
         **({'markup_override': '1.03'} if mode == 'detached' else {}),
         'materials': [{'item': '物料', 'usage': 1, 'unit_price_hkd': '.25885',
                        'unit_price_source_currency': 'HKD', 'markup': 1}]}
        for category in ('clothes', 'hair')
    ]}}
    workbook, context, summary = export_costs(payloads, mode)
    try:
        assert context['factory_price_hkd'] == Decimal('.5177')
        entries = summary['shipping_pricing']['pricing_entries']
        assert sum(Decimal(e['amount_hkd']) for e in entries) == Decimal('.5177')
        assert sum(Decimal(g['cost_hkd']) for g in summary['shipping_pricing']['pricing_groups']) == Decimal('.5177')
        sheet = workbook['报价明细']
        assert sum(row_amounts(sheet, '车衣') + row_amounts(sheet, '车发')) == pytest.approx(.5177)
        assert not any(cell.value == '车缝未分类成本' for row in sheet for cell in row)
    finally:
        workbook.close()
