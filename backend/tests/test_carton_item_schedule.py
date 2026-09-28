from datetime import datetime
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from app.services.carton_procurement_imports import _parse_weekly, _match_rows


def _prepare_schedule_customer(client):
    from test_molding_sample_api import login_as
    from test_carton_procurement_api import _ensure_dickie_customer
    login_as(client, 'admin')
    _ensure_dickie_customer(client)


HEADERS = ['订单类型', 'Contract No.', 'SO#/Reference', 'P/O#:', '客名', '产品编号',
           '产品中文名称', '數量', '装箱', '验货日期', '客要求走货期']


def workbook(rows, *, second_rows=None, reverse=False):
    book = Workbook()
    book.active.title = '接单表'
    book.active.append(['合同号', '货号', '数量'])
    book.active.append(['WRONG', 'IGNORE', 999])
    for name, values in [('子弹枪ITEM表', rows), ('水枪ITEM表', second_rows)]:
        if values is None:
            continue
        sheet = book.create_sheet(name)
        sheet.append(['模板说明，不是业务数据'])
        sheet.append(['合同基本信息'])
        sheet.append(HEADERS[::-1] if reverse else HEADERS)
        for row in values:
            sheet.append(row[::-1] if reverse else row)
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def source_row(kind='正单', contract='53135', item='0040220-1', qty=360):
    return [kind, contract, '', '0989551378', 'WMU', item, '双支三弹枪', qty, 6,
            datetime(2026, 9, 10), datetime(2026, 9, 17)]


def test_item_sheets_only_by_header_preserve_po_and_nonformal_rows():
    result = _parse_weekly('schedule.xlsx', workbook(
        [source_row(), source_row('备料单', '', qty=0), source_row('正式PO', '53136')],
        second_rows=[source_row(contract='53137')], reverse=True))
    assert len(result['rows']) == 4
    first = result['rows'][0]
    assert first['contract_no'] == first['reference'] == '53135'
    assert first['source_reference'] == ''
    assert first['customer_po'] == '0989551378'
    assert first['item_no'] == '0040220-1'
    assert first['customer_due_date'] == '2026-09-17'
    assert first['inspection_window'] == '2026-09-10'
    assert first['source_row'] == 4
    assert result['review_count'] == 2
    assert result['rows'][1]['quantity'] == 0
    assert result['rows'][1]['contract_no'] == ''
    assert result['field_mappings'][0]['fields']['quantity'] == '數量'


def test_duplicate_formal_rows_and_missing_keys_require_review():
    result = _parse_weekly('schedule.xlsx', workbook([
        source_row(), source_row(), source_row(contract=''),
        source_row(contract='53136', qty=0), source_row(contract='53137', qty=None)]))
    assert all(row['match_status'] == 'REVIEW_REQUIRED' for row in result['rows'])
    assert result['review_count'] == 5


def test_item_rows_not_silently_truncated_at_old_preview_limit():
    result = _parse_weekly('schedule.xlsx', workbook([source_row(contract=str(i)) for i in range(501)]))
    assert len(result['rows']) == 501


def test_missing_required_header_is_explicit_error():
    book = Workbook()
    book.active.title = 'ITEM表'
    book.active.append(['产品编号', '數量'])
    stream = BytesIO()
    book.save(stream)
    with pytest.raises(HTTPException) as error:
        _parse_weekly('schedule.xlsx', stream.getvalue())
    assert error.value.status_code == 422
    assert '缺少' in error.value.detail


def test_duplicate_quantity_headers_are_rejected_instead_of_guessing():
    book = Workbook()
    book.active.title = 'ITEM表'
    book.active.append(HEADERS + ['数量'])
    book.active.append(source_row() + [999])
    stream = BytesIO()
    book.save(stream)
    with pytest.raises(HTTPException) as error:
        _parse_weekly('schedule.xlsx', stream.getvalue())
    assert '重复' in error.value.detail


def test_oversized_item_sheet_fails_instead_of_truncating_demand():
    book = Workbook()
    book.active.title = 'ITEM表'
    book.active.append(HEADERS)
    book.active.cell(5001, 1, '正单')
    stream = BytesIO()
    book.save(stream)
    with pytest.raises(HTTPException) as error:
        _parse_weekly('schedule.xlsx', stream.getvalue())
    assert error.value.status_code == 422
    assert '超过' in error.value.detail


def test_review_rows_do_not_get_automatic_inspection_delivery_reminders():
    from datetime import date
    from app.services.carton_procurement_imports import _apply_inspection_reminders
    rows = _parse_weekly('schedule.xlsx', workbook([source_row('样板单')]))['rows']
    _apply_inspection_reminders(rows, advance_days=3, reference_date=date(2026, 9, 21))
    assert rows[0]['reminder_status'] == 'REVIEW_REQUIRED'
    assert '不计入正单' in rows[0]['suggestion']


class MatchingDb:
    def __init__(self, orders):
        self.orders = orders

    def execute(self, query):
        return SimpleNamespace(all=lambda: [(SimpleNamespace(), order) for order in self.orders])

    def scalars(self, query):
        return SimpleNamespace(all=lambda: [])


def order(**updates):
    fields = dict(factory_id='huaxing', id='o1', order_no='CO1', customer_po='0989551378', contract_no='53135',
                  item_no='0040220-1', customer_code='BUZZBEE', customer_name='BUZZ BEE',
                  product_order_quantity=360, customer_due_date='2026-09-17', status='PENDING_SUPPLIER')
    return SimpleNamespace(**(fields | updates))


def test_unified_matching_never_falls_back_to_only_contract_or_only_item():
    rows = _parse_weekly('schedule.xlsx', workbook([source_row()]))['rows']
    _match_rows(MatchingDb([order(contract_no='OTHER')]), 'huaxing', 'WEEKLY_SCHEDULE', rows)
    assert rows[0]['match_status'] == 'MISSING_ORDER'
    assert rows[0]['procurement_state'] == 'NEEDS_ORDER'
    _match_rows(MatchingDb([order(item_no='OTHER')]), 'huaxing', 'WEEKLY_SCHEDULE', rows)
    assert rows[0]['match_status'] == 'MISSING_ORDER'


@pytest.mark.parametrize('status,quantity,state', [
    ('CONFIRMED', 360, 'NEEDS_ORDER'), ('PENDING_SUPPLIER', 360, 'ORDERED'),
    ('COMPLETED', 360, 'COMPLETED'), ('COMPLETED', 350, 'NEEDS_ORDER'),
])
def test_procurement_status_and_quantity_gap(status, quantity, state):
    rows = _parse_weekly('schedule.xlsx', workbook([source_row()]))['rows']
    _match_rows(MatchingDb([order(status=status, product_order_quantity=quantity)]), 'huaxing', 'WEEKLY_SCHEDULE', rows)
    assert rows[0]['procurement_state'] == state
    assert rows[0]['match_status'] == ('MATCHED' if quantity == 360 else 'QUANTITY_MISMATCH')


def test_customer_deadline_difference_and_nonformal_exclusion():
    rows = _parse_weekly('schedule.xlsx', workbook([source_row(), source_row('样板单')]))['rows']
    _match_rows(MatchingDb([order(customer_due_date='2026-09-15')]), 'huaxing', 'WEEKLY_SCHEDULE', rows)
    assert rows[0]['match_status'] == 'DATE_MISMATCH'
    assert rows[0]['date_difference']['business_date'] == '2026-09-17'
    assert rows[1]['match_status'] == 'REVIEW_REQUIRED'
    assert 'order_id' not in rows[1]


@pytest.mark.parametrize('raw', ['9/1-10% 9/14-100%', '9/21', '2026/9/1-10% 2026/9/14-100%', '2026-02-30'])
@pytest.mark.parametrize('column,field,source', [
    (9, 'inspection_window', 'source_inspection_window'),
    (10, 'customer_due_date', 'source_customer_due_date'),
])
def test_uncertain_item_dates_preserve_evidence_without_guessing(raw, column, field, source):
    from datetime import date
    from app.services.carton_procurement_imports import _apply_inspection_reminders
    values = source_row()
    values[column] = raw
    parsed = _parse_weekly('schedule.xlsx', workbook([values]))
    row = parsed['rows'][0]
    assert row[source] == raw
    assert row[field] == ''
    assert row['date_review_required'] is True
    assert parsed['review_count'] == 1
    _match_rows(MatchingDb([order()]), 'huaxing', 'WEEKLY_SCHEDULE', [row])
    assert row['order_id'] == 'o1'
    assert row['procurement_state'] == 'ORDERED'
    assert row['match_status'] == 'REVIEW_REQUIRED'
    _apply_inspection_reminders([row], advance_days=3, reference_date=date(2026, 9, 21))
    assert row['reminder_status'] == 'REVIEW_REQUIRED'
    assert not row.get('required_delivery_date')
    assert row[source] == raw


def test_import_api_persists_review_and_is_idempotent_without_writing_orders(monkeypatch):
    from test_molding_sample_api import make_client, login_as
    with make_client(monkeypatch) as client:
        _prepare_schedule_customer(client)
        login_as(client, 'warehouse_keeper')
        content = workbook([source_row(), source_row('样板单')])
        kwargs = dict(params={'factory_id': 'huaxing'}, files={'file': ('unified.xlsx', content)})
        response = client.post('/api/carton-procurement/weekly-imports', data={'customer_code': 'DICKIE'}, **kwargs)
        assert response.status_code == 201, response.text
        body = response.json()
        assert body['parse_summary']['review_count'] == 1
        assert [row['match_status'] for row in body['parse_summary']['rows']] == ['MISSING_ORDER', 'REVIEW_REQUIRED']
        repeat = client.post('/api/carton-procurement/weekly-imports', data={'customer_code': 'DICKIE'}, **kwargs)
        assert repeat.status_code == 201, repeat.text
        assert repeat.json()['id'] == body['id']
        assert repeat.json()['duplicate'] is True
        exceptions = client.get('/api/carton-procurement/exceptions', params={'factory_id': 'huaxing'}).json()
        assert {row['category'] for row in exceptions['items']} == {'MISSING_ORDER', 'SCHEDULE_REVIEW_REQUIRED'}
        assert exceptions['total'] == 2
        assert client.get('/api/carton-procurement/orders', params={'factory_id': 'huaxing'}).json()['total'] == 0
        assert client.get('/api/carton-procurement/receipts', params={'factory_id': 'huaxing'}).json()['total'] == 0


def test_item_section_titles_override_row_order_type_and_shipped_is_not_pending():
    parsed = _parse_weekly('schedule.xlsx', workbook([
        source_row(), ['取消单'], source_row(contract='53136'),
        ['已走货订单'], source_row(contract='53137'),
    ]))
    assert [row['schedule_section'] for row in parsed['rows']] == ['PENDING', 'CANCELLED', 'SHIPPED']
    assert [row['source_row'] for row in parsed['rows']] == [4, 6, 8]
    assert parsed['rows'][1]['match_status'] == 'REVIEW_REQUIRED'
    assert parsed['rows'][2]['match_status'] == 'REVIEW_REQUIRED'


def test_schedule_comparison_only_flags_explicit_new_and_cancelled_transition():
    from app.services.carton_schedule_tracking import annotate_changes, customer_key, identity

    previous = [
        dict(source_customer_name='WMU', contract_no='53135', customer_po='0989551378',
             item_no='0040220-1', order_type='正单', schedule_section='PENDING'),
        dict(source_customer_name='WMU', contract_no='53138', customer_po='',
             item_no='0040220-2', order_type='正单', schedule_section='PENDING'),
    ]
    rows = [
        dict(previous[0], customer_po='PO-LATER-FILLED', schedule_section='CANCELLED'),
        dict(previous[0], contract_no='53136', schedule_section='PENDING'),
    ]
    annotate_changes(rows, {customer_key(previous[0]): previous}, {identity(previous[0])})
    assert [row['schedule_change'] for row in rows] == ['CANCELLED_AFTER_ORDER', 'NEW']
    assert rows[0]['manual_ordered'] is True
    assert all(row['schedule_change'] != 'CANCELLED' for row in rows[1:])
    duplicates = [dict(previous[0], customer_po='PO-A'), dict(previous[0], customer_po='PO-B')]
    annotate_changes(duplicates, {}, set())
    assert [row['schedule_change'] for row in duplicates] == ['REVIEW_REQUIRED', 'REVIEW_REQUIRED']


def test_manual_order_mark_survives_import_and_cancelled_change_is_actionable(monkeypatch):
    from test_molding_sample_api import make_client, login_as
    with make_client(monkeypatch) as client:
        _prepare_schedule_customer(client)
        login_as(client, 'warehouse_keeper')
        baseline = client.post('/api/carton-procurement/weekly-imports', data={'customer_code': 'DICKIE'},
            params={'factory_id': 'huaxing'}, files={'file': ('baseline.xlsx', workbook([source_row()]))})
        assert baseline.status_code == 201, baseline.text
        batch_id = baseline.json()['id']
        assert baseline.json()['parse_summary']['rows'][0]['schedule_change'] == 'BASELINE'
        mark = client.post('/api/carton-procurement/schedule-order-marks', json={
            'factory_id': 'huaxing', 'batch_id': batch_id, 'source_sheet': '子弹枪ITEM表',
            'source_row': 4, 'marked': True,
        })
        assert mark.status_code == 200, mark.text
        key = mark.json()['identity']
        marks = client.get('/api/carton-procurement/schedule-order-marks', params={'factory_id': 'huaxing'})
        assert marks.json()[key]['marked'] is True
        changed = client.post('/api/carton-procurement/weekly-imports', data={'customer_code': 'DICKIE'},
            params={'factory_id': 'huaxing'}, files={'file': ('changed.xlsx', workbook([
                source_row(contract='53136'), ['取消单'], source_row(),
            ]))})
        assert changed.status_code == 201, changed.text
        rows = changed.json()['parse_summary']['rows']
        assert [row['schedule_change'] for row in rows] == ['NEW', 'CANCELLED_AFTER_ORDER']
        exceptions = client.get('/api/carton-procurement/exceptions', params={'factory_id': 'huaxing'}).json()['items']
        categories = {row['category'] for row in exceptions if row['source_id'] == changed.json()['id']}
        assert categories == {'SCHEDULE_NEW_ORDER', 'SCHEDULE_CANCELLED_AFTER_ORDER'}
        assert client.get('/api/carton-procurement/orders', params={'factory_id': 'huaxing'}).json()['total'] == 0


def test_bulk_marks_are_atomic_idempotent_and_keep_ordered_cancellation_reminders(monkeypatch):
    from test_molding_sample_api import make_client, login_as
    with make_client(monkeypatch) as client:
        _prepare_schedule_customer(client)
        login_as(client, 'warehouse_keeper')
        baseline = client.post('/api/carton-procurement/weekly-imports', data={'customer_code': 'DICKIE'}, params={'factory_id': 'huaxing'},
            files={'file': ('bulk.xlsx', workbook([source_row(), source_row(contract='53136'),
                                                  ['取消单'], source_row(contract='53137')]))})
        assert baseline.status_code == 201, baseline.text
        payload = {'factory_id': 'huaxing', 'batch_id': baseline.json()['id'], 'marked': True,
                   'rows': [{'source_sheet': '子弹枪ITEM表', 'source_row': 4},
                            {'source_sheet': '子弹枪ITEM表', 'source_row': 7}]}
        rejected = client.post('/api/carton-procurement/schedule-order-marks/bulk', json=payload)
        assert rejected.status_code == 422, rejected.text
        assert client.get('/api/carton-procurement/schedule-order-marks', params={'factory_id': 'huaxing'}).json() == {}
        payload['rows'][1]['source_row'] = 5
        marked = client.post('/api/carton-procurement/schedule-order-marks/bulk', json=payload)
        assert marked.status_code == 200, marked.text
        assert marked.json()['changed_count'] == 2
        keys = {row['identity'] for row in marked.json()['items']}
        repeat = client.post('/api/carton-procurement/schedule-order-marks/bulk', json=payload)
        assert repeat.status_code == 200 and repeat.json()['changed_count'] == 0
        events = client.get('/api/carton-procurement/audit-events', params={'factory_id': 'huaxing'}).json()['items']
        marks = [event for event in events if event['event_type'] == 'SCHEDULE_ORDER_MARK_CHANGED']
        assert len(marks) == 2
        assert len({event['detail']['operation_id'] for event in marks}) == 1
        assert {event['entity_id'] for event in marks} == keys
        changed = client.post('/api/carton-procurement/weekly-imports', data={'customer_code': 'DICKIE'}, params={'factory_id': 'huaxing'},
            files={'file': ('next.xlsx', workbook([source_row(contract='53136'), source_row(contract='53138'),
                                                  ['取消单'], source_row()]))})
        assert changed.status_code == 201, changed.text
        rows = changed.json()['parse_summary']['rows']
        assert rows[0]['manual_ordered'] is True and rows[2]['schedule_change'] == 'CANCELLED_AFTER_ORDER'
        exceptions = client.get('/api/carton-procurement/exceptions', params={'factory_id': 'huaxing'}).json()['items']
        new_exceptions = [row for row in exceptions if row['source_id'] == changed.json()['id']]
        assert {row['category'] for row in new_exceptions} == {'SCHEDULE_NEW_ORDER', 'SCHEDULE_CANCELLED_AFTER_ORDER'}
        assert all(row['contract_no'] != '53136' for row in new_exceptions)
        assert client.get('/api/carton-procurement/orders', params={'factory_id': 'huaxing'}).json()['total'] == 0


def test_bulk_marks_reject_duplicate_rows_identity_collisions_wrong_factory_and_permission(monkeypatch):
    from test_molding_sample_api import make_client, login_as
    with make_client(monkeypatch) as client:
        _prepare_schedule_customer(client)
        login_as(client, 'admin')
        baseline = client.post('/api/carton-procurement/weekly-imports', data={'customer_code': 'DICKIE'}, params={'factory_id': 'huaxing'},
            files={'file': ('duplicates.xlsx', workbook([source_row(), source_row(), source_row(contract='53136')]))})
        assert baseline.status_code == 201, baseline.text
        source = {'source_sheet': '子弹枪ITEM表', 'source_row': 6}
        payload = {'factory_id': 'huaxing', 'batch_id': baseline.json()['id'], 'marked': True, 'rows': [source]}
        assert client.post('/api/carton-procurement/schedule-order-marks/bulk', json={**payload, 'rows': [source, source]}).status_code == 422
        assert client.post('/api/carton-procurement/schedule-order-marks/bulk', json={**payload, 'rows': [source] * 101}).status_code == 422
        assert client.post('/api/carton-procurement/schedule-order-marks/bulk', json={**payload, 'rows': [source, {'source_sheet': '子弹枪ITEM表', 'source_row': 4}]}).status_code == 409
        assert client.post('/api/carton-procurement/schedule-order-marks/bulk', json={**payload, 'factory_id': 'huadeng'}).status_code == 404
        assert client.get('/api/carton-procurement/schedule-order-marks', params={'factory_id': 'huaxing'}).json() == {}
        login_as(client, 'engineer')
        denied = client.post('/api/carton-procurement/schedule-order-marks/bulk', json=payload)
        assert denied.status_code == 403, denied.text


def test_previously_marked_order_in_first_cancelled_snapshot_is_not_silenced():
    from app.services.carton_schedule_tracking import annotate_changes, identity
    cancelled = dict(source_customer_name='WMU', contract_no='53135', item_no='0040220-1',
                     order_type='正单', schedule_section='CANCELLED')
    annotate_changes([cancelled], {}, {identity(cancelled)})
    assert cancelled['schedule_change'] == 'CANCELLED_AFTER_ORDER'


def test_source_order_link_survives_customer_selection_reimport_and_cancellation(monkeypatch):
    from test_molding_sample_api import make_client, login_as
    from test_carton_procurement_api import _ensure_dickie_customer, _order_payload, _submit_order
    with make_client(monkeypatch) as client:
        _prepare_schedule_customer(client)
        login_as(client, 'admin')
        _ensure_dickie_customer(client)
        values = source_row()
        values[3] = ''  # Explicit source linkage also protects legacy orders with no PO.
        batch = client.post('/api/carton-procurement/weekly-imports', data={'customer_code': 'DICKIE'}, params={'factory_id': 'huaxing'},
            files={'file': ('source.xlsx', workbook([values]))})
        assert batch.status_code == 201, batch.text
        row = batch.json()['parse_summary']['rows'][0]
        payload = {**_order_payload(), 'contract_no': '53135', 'item_no': '0040220-1',
                   'product_order_quantity': 360, 'customer_due_date': '2026-09-17',
                   'schedule_source': {'batch_id': batch.json()['id'], 'source_sheet': '子弹枪ITEM表', 'source_row': 4}}
        created = client.post('/api/carton-procurement/orders', json=payload)
        assert created.status_code == 201, created.text
        saved = created.json()
        assert saved['status'] == 'CONFIRMED' and saved['customer_code'] == 'DICKIE'
        state = client.get('/api/carton-procurement/schedule-order-marks', params={'factory_id': 'huaxing'}).json()
        assert state[row['schedule_identity']]['order_ids'] == [saved['id']]
        assert state[row['schedule_identity']]['marked'] is False
        assert client.post('/api/carton-procurement/orders', json=payload).status_code == 409
        assert client.get('/api/carton-procurement/orders', params={'factory_id': 'huaxing'}).json()['total'] == 1
        _submit_order(client, saved)
        changed_values = list(values)
        changed_values[7] = 361
        changed = client.post('/api/carton-procurement/weekly-imports', data={'customer_code': 'DICKIE'}, params={'factory_id': 'huaxing'},
            files={'file': ('quantity-change.xlsx', workbook([changed_values]))})
        assert changed.status_code == 201, changed.text
        linked = changed.json()['parse_summary']['rows'][0]
        assert linked['source_customer_name'] == 'WMU' and linked['customer_code'] == 'DICKIE'
        assert linked['order_id'] == saved['id'] and linked['match_status'] == 'QUANTITY_MISMATCH'
        cancelled_values = list(values)
        cancelled_values[0] = ''
        cancelled = client.post('/api/carton-procurement/weekly-imports', data={'customer_code': 'DICKIE'}, params={'factory_id': 'huaxing'},
            files={'file': ('cancelled.xlsx', workbook([['取消单'], cancelled_values]))})
        assert cancelled.status_code == 201, cancelled.text
        cancelled_row = cancelled.json()['parse_summary']['rows'][0]
        assert cancelled_row['order_id'] == saved['id']
        assert cancelled_row['schedule_change'] == 'CANCELLED_AFTER_ORDER'
        exceptions = client.get('/api/carton-procurement/exceptions', params={'factory_id': 'huaxing'}).json()['items']
        assert any(item['source_id'] == cancelled.json()['id'] and item['category'] == 'SCHEDULE_CANCELLED_AFTER_ORDER'
                   and item['severity'] == 'HIGH' for item in exceptions)
        orders = client.get('/api/carton-procurement/orders', params={'factory_id': 'huaxing'}).json()['items']
        assert orders[0]['status'] == 'PENDING_SUPPLIER'  # Reminder does not cancel a purchase.


def test_source_order_creation_rejects_marked_ambiguous_stale_and_other_factory_sources(monkeypatch):
    from test_molding_sample_api import make_client, login_as
    from test_carton_procurement_api import _ensure_dickie_customer, _order_payload
    with make_client(monkeypatch) as client:
        _prepare_schedule_customer(client)
        login_as(client, 'admin')
        _ensure_dickie_customer(client)
        login_as(client, 'admin')
        batch = client.post('/api/carton-procurement/weekly-imports', data={'customer_code': 'DICKIE'}, params={'factory_id': 'huaxing'},
            files={'file': ('baseline.xlsx', workbook([source_row(), source_row(contract='53136')]))})
        assert batch.status_code == 201, batch.text
        source = {'batch_id': batch.json()['id'], 'source_sheet': '子弹枪ITEM表', 'source_row': 4}
        payload = {**_order_payload(), 'contract_no': '53135', 'item_no': '0040220-1', 'schedule_source': source}
        assert client.post('/api/carton-procurement/orders', json={**payload, 'factory_id': 'huadeng'}).status_code == 404
        assert client.post('/api/carton-procurement/orders', json={**payload, 'item_no': 'OTHER'}).status_code == 422
        marked = client.post('/api/carton-procurement/schedule-order-marks', json={
            'factory_id': 'huaxing', **source, 'marked': True})
        assert marked.status_code == 200, marked.text
        assert client.post('/api/carton-procurement/orders', json=payload).status_code == 409
        client.post('/api/carton-procurement/weekly-imports', data={'customer_code': 'DICKIE'}, params={'factory_id': 'huaxing'},
            files={'file': ('next.xlsx', workbook([['取消单'], source_row(contract='53136')]))})
        stale = {**payload, 'contract_no': '53136', 'schedule_source': {**source, 'source_row': 5}}
        assert client.post('/api/carton-procurement/orders', json=stale).status_code == 409
        duplicate = client.post('/api/carton-procurement/weekly-imports', data={'customer_code': 'DICKIE'}, params={'factory_id': 'huaxing'},
            files={'file': ('duplicate.xlsx', workbook([source_row(contract='53137'), source_row(contract='53137')]))})
        assert duplicate.status_code == 201, duplicate.text
        ambiguous = {**payload, 'contract_no': '53137', 'schedule_source': {**source, 'batch_id': duplicate.json()['id']}}
        assert client.post('/api/carton-procurement/orders', json=ambiguous).status_code == 409
        assert client.get('/api/carton-procurement/orders', params={'factory_id': 'huaxing'}).json()['total'] == 0


def test_weekly_import_requires_active_factory_customer_and_isolated_customer_identity(monkeypatch):
    from test_molding_sample_api import make_client, login_as
    from test_carton_procurement_api import _order_payload
    with make_client(monkeypatch) as client:
        _prepare_schedule_customer(client)
        login_as(client, 'admin')
        for factory, code, status in [('huaxing', 'OTHER', 'ACTIVE'), ('huaxing', 'DISABLED', 'INACTIVE'),
                                      ('huadeng', 'OTHER-FACTORY', 'ACTIVE')]:
            customer = client.post('/api/carton-procurement/customers', json={
                'factory_id': factory, 'customer_code': code, 'customer_name': code, 'status': status})
            assert customer.status_code == 201, customer.text
        content = workbook([source_row()])

        def upload(code=None, data=content, filename='business.xlsx'):
            return client.post('/api/carton-procurement/weekly-imports', params={'factory_id': 'huaxing'},
                data={'customer_code': code} if code is not None else {}, files={'file': (filename, data)})

        assert upload().status_code == 422
        for invalid in ['', 'NOT-FOUND', 'DISABLED', 'OTHER-FACTORY']:
            assert upload(invalid).status_code == 422
        assert client.get('/api/carton-procurement/imports', params={'factory_id': 'huaxing', 'import_type': 'WEEKLY_SCHEDULE'}).json()['total'] == 0
        first = upload('DICKIE')
        assert first.status_code == 201, first.text
        first_row = first.json()['parse_summary']['rows'][0]
        assert first_row['source_customer_name'] == 'WMU'
        assert first_row['customer_name'] == 'Dickie' and first_row['schedule_customer_code'] == 'DICKIE'
        source = {'batch_id': first.json()['id'], 'source_sheet': '子弹枪ITEM表', 'source_row': 4}
        wrong_customer = {**_order_payload(), 'contract_no': '53135', 'item_no': '0040220-1',
                          'customer_code': 'OTHER', 'customer_name': 'OTHER', 'schedule_source': source}
        assert client.post('/api/carton-procurement/orders', json=wrong_customer).status_code == 422
        created = client.post('/api/carton-procurement/orders', json={**_order_payload(),
            'contract_no': '53135', 'item_no': '0040220-1', 'customer_po': '0989551378', 'product_order_quantity': 360})
        assert created.status_code == 201, created.text
        other = upload('OTHER')
        assert other.status_code == 201, other.text
        other_row = other.json()['parse_summary']['rows'][0]
        assert other.json()['id'] != first.json()['id']
        assert other_row['schedule_identity'] != first_row['schedule_identity']
        assert other_row['match_status'] == 'MISSING_ORDER' and not other_row.get('order_id')
        marked = client.post('/api/carton-procurement/schedule-order-marks', json={
            'factory_id': 'huaxing', 'batch_id': other.json()['id'], 'source_sheet': '子弹枪ITEM表', 'source_row': 4, 'marked': True})
        assert marked.status_code == 200, marked.text
        renamed_buyer = source_row()
        renamed_buyer[4] = 'ABU ISSA'
        renamed_buyer[3] = 'FILLED-PO'
        next_import = upload('OTHER', workbook([renamed_buyer]), 'later.xlsx')
        assert next_import.status_code == 201, next_import.text
        next_row = next_import.json()['parse_summary']['rows'][0]
        assert next_row['schedule_identity'] == other_row['schedule_identity'] and next_row['manual_ordered'] is True
        cancelled = upload('OTHER', workbook([['取消单'], renamed_buyer]), 'cancelled.xlsx')
        assert cancelled.status_code == 201, cancelled.text
        assert cancelled.json()['parse_summary']['rows'][0]['schedule_change'] == 'CANCELLED_AFTER_ORDER'
        repeat = upload('DICKIE')
        assert repeat.status_code == 201 and repeat.json()['duplicate'] is True
        assert repeat.json()['id'] == first.json()['id']
        assert client.get('/api/carton-procurement/orders', params={'factory_id': 'huaxing'}).json()['total'] == 1
