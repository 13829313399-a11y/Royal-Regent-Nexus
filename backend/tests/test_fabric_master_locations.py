import importlib
import json
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4, uuid5
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session
from test_fabric_chase_contracts import migration


@pytest.fixture
def setup(tmp_path):
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    engine = create_engine('sqlite:///' + str(tmp_path / 'locations.db'))
    with engine.begin() as conn, Operations.context(MigrationContext.configure(conn)):
        for name in ('20261005_0134_fabric_procurement.py', '20261008_0142_fabric_receiving.py', '20261008_0143_fabric_master_chase.py'):
            migration(name).upgrade()
    with Session(engine) as db:
        yield db, SimpleNamespace(id='A', display_name='负责人'), importlib.import_module('app.services.fabric_master_locations')
    engine.dispose()


def body(rows):
    from app.schemas.fabric_master import LocationPreviewRequest
    return LocationPreviewRequest(factory_id='huakang-c', rows=rows)


def apply_body(rows, token, request_id=None):
    from app.schemas.fabric_master import LocationApplyRequest
    return LocationApplyRequest(factory_id='huakang-c', rows=rows, preview_token=token, request_id=request_id or uuid4())


def add(db, actor, service, warehouse='布料仓', bins='A01-A03'):
    payload = body([{'warehouse': warehouse, 'bins': bins}])
    plan = service.preview(db, actor, payload)
    request = apply_body(payload.rows, plan['preview_token'])
    result = service.apply(db, actor, request)
    return request, result


def test_categorized_positions_conflicts_and_legacy_request_signatures(setup):
    db, actor, service = setup
    rows = [{'warehouse': '原料仓', 'bins': '布料1架、布料12架、布料1区', 'category': 'FABRIC'},
            {'warehouse': '原料仓', 'bins': '辅料1架、辅料1区', 'category': 'ACCESSORY'},
            {'warehouse': '原料仓', 'bins': '配线1区、配线5区', 'category': 'THREAD'}]
    preview = service.preview(db, actor, body(rows))
    assert not preview['errors'] and preview['new'] == 7
    request = apply_body(rows, preview['preview_token'])
    result = service.apply(db, actor, request)
    assert service.apply(db, actor, request) == result
    from app.services.warehouse_locations import catalog, resolve
    places = {row['code']: row for row in catalog(db, 'fabric')}
    assert places['配线5区']['category'] == 'THREAD'
    assert resolve(db, 'fabric', places['布料12架']['id'], material_category='FABRIC')['category'] == 'FABRIC'
    with pytest.raises(HTTPException):
        resolve(db, 'fabric', places['布料12架']['id'], material_category='THREAD')
    conflict = service.preview(db, actor, body([{'warehouse': '原料仓', 'bins': '布料1架、布料2架', 'category': 'THREAD'}]))
    assert conflict['errors']
    with pytest.raises(HTTPException):
        service.apply(db, actor, apply_body([{'warehouse': '原料仓', 'bins': '布料1架、布料2架', 'category': 'THREAD'}], conflict['preview_token']))
    db.rollback()
    assert len(service.locations(db)) == 7
    legacy, result = add(db, actor, service)
    from app.models.fabric_master import FabricMasterChange
    from app.services.fabric_procurement_parser import digest
    old_body = legacy.model_dump(mode='json')
    for row in old_body['rows']:
        row.pop('category')
    change = db.query(FabricMasterChange).filter_by(request_id=str(legacy.request_id)).one()
    assert change.request_hash == digest([actor.id, old_body])
    assert service.apply(db, actor, legacy) == result
    # Uncategorized legacy entries remain usable for all raw-material categories.
    general = next(row for row in catalog(db, 'fabric') if row['code'] == 'A01')
    assert resolve(db, 'fabric', general['id'], material_category='THREAD')


def rename_body(group, new_name='新布料仓'):
    from app.schemas.fabric_master import WarehouseRenameRequest
    return WarehouseRenameRequest(factory_id='huakang-c', request_id=uuid4(), warehouse='布料仓', name=new_name,
                                  expected_group_token=importlib.import_module('app.services.fabric_master').warehouse_tokens(group)['布料仓'])


def manage_body(rows, action='DELETE', category=''):
    from app.schemas.fabric_master import LocationManageRequest
    return LocationManageRequest(factory_id='huakang-c', request_id=uuid4(), action=action, category=category,
        selections=[{'id': row.id, 'expected_revision': row.revision} for row in rows])


def test_batch_category_updates_only_selected_existing_positions_and_checks_all_revisions(setup):
    db, actor, service = setup
    add(db, actor, service)
    rows = service.locations(db)
    stale = manage_body(rows, 'CATEGORY', 'THREAD')
    result = service.manage(db, actor, manage_body(rows[:1], 'CATEGORY', 'FABRIC'))
    assert result == {'action': 'CATEGORY', 'updated': 1, 'stock_posted': False}
    with pytest.raises(HTTPException) as error:
        service.manage(db, actor, stale)
    assert error.value.status_code == 409
    db.rollback()
    assert len(service.locations(db)) == 3
    assert [json.loads(row.data_json)['category'] for row in service.locations(db)] == ['FABRIC', '', '']
    for table in ('fabric_receipts', 'fabric_stock_batches', 'fabric_inventory_movements'):
        assert db.scalar(text(f'SELECT count(*) FROM {table}')) == 0


def test_delete_restore_reserves_identity_preserves_status_and_replays_without_redeleting(setup):
    from app.schemas.fabric_master import SaveMasterRequest
    from app.services.warehouse_locations import catalog, resolve
    db, actor, service = setup
    add(db, actor, service)
    rows = service.locations(db)
    rows[1].status = 'INACTIVE'; db.commit()
    request = manage_body(rows)
    result = service.manage(db, actor, request)
    assert result['updated'] == 3
    assert not service.locations(db) and not catalog(db, 'fabric')
    listed = service.master.list_records(db, 'LOCATION')
    assert not listed['items'] and not listed['warehouse_tokens'] and len(listed['deleted_items']) == 3
    assert len(service.master.history(db, rows[0].id)['items']) == 2
    assert service.preview(db, actor, body([{'warehouse': '新仓', 'bins': 'a01'}]))['errors']
    with pytest.raises(HTTPException):
        resolve(db, 'fabric', rows[0].id)
    with pytest.raises(HTTPException) as error:
        service.master.save(db, actor, SaveMasterRequest(factory_id='huakang-c', request_id=uuid4(), kind='LOCATION',
            id=rows[0].id, expected_revision=rows[0].revision, code=rows[0].code, name=rows[0].name,
            status='ACTIVE', data={'warehouse': '布料仓'}))
    assert error.value.status_code == 409
    db.rollback()
    assert service.manage(db, actor, manage_body(service.locations(db, include_deleted=True), 'RESTORE'))['updated'] == 3
    assert [row.status for row in service.locations(db)] == ['ACTIVE', 'INACTIVE', 'ACTIVE']
    assert service.manage(db, actor, request) == result
    assert len(service.locations(db)) == 3
    with pytest.raises(HTTPException) as error:
        service.manage(db, SimpleNamespace(id='other', display_name='另一个人'), request)
    assert error.value.status_code == 409
    db.rollback()
    with pytest.raises(HTTPException):
        service.manage(db, actor, request.model_copy(update={'action': 'RESTORE'}))


@pytest.mark.parametrize('reference', ['stable-id', 'historical-label', 'legacy-code', 'semi-only'])
def test_delete_rejects_referenced_history_after_warehouse_rename_and_never_partially_deletes(setup, reference):
    from app.models.warehouse_operations import WarehouseOperation
    db, actor, service = setup
    WarehouseOperation.__table__.create(db.get_bind())
    add(db, actor, service)
    rows = service.locations(db)
    facts = {'stable-id': {'destination': {'location_id': rows[0].id}},
             'historical-label': {'location': '布料仓／A01'},
             'legacy-code': {'location': 'a01'},
             'semi-only': {'location': 'A01'}}[reference]
    db.add(WarehouseOperation(id=uuid4().hex, factory_id='huakang-c', warehouse='semi' if reference == 'semi-only' else 'fabric',
        sequence=1, request_id=uuid4().hex, request_hash='hash', kind='OPENING', business_date='2026-10-13',
        payload_json=json.dumps(facts), actor_id=actor.id, actor_name=actor.display_name, occurred_at='2026-10-13T12:00:00+08:00'))
    db.commit()
    service.rename(db, actor, rename_body(service.locations(db)))
    request = manage_body(service.locations(db))
    if reference == 'semi-only':
        assert service.manage(db, actor, request)['updated'] == 3
        return
    with pytest.raises(HTTPException) as error:
        service.manage(db, actor, request)
    assert error.value.status_code == 409 and 'A01' in error.value.detail
    db.rollback()
    assert len(service.locations(db)) == 3
    assert all(row.revision == 2 for row in service.locations(db))


def test_management_write_failure_rolls_back_all_tombstones_and_audit(setup, monkeypatch):
    db, actor, service = setup
    add(db, actor, service)
    request = manage_body(service.locations(db))
    original_flush = db.flush
    writes = 0
    def fail_second(*args, **kwargs):
        nonlocal writes
        if db.dirty:
            writes += 1
            if writes == 2:
                raise RuntimeError('simulated management failure')
        return original_flush(*args, **kwargs)
    monkeypatch.setattr(db, 'flush', fail_second)
    with pytest.raises(RuntimeError):
        service.manage(db, actor, request)
    db.rollback()
    monkeypatch.setattr(db, 'flush', original_flush)
    assert len(service.locations(db)) == 3
    assert all(row.revision == 1 for row in service.locations(db))
    assert db.scalar(text('SELECT count(*) FROM fabric_master_changes')) == 3


def test_batch_ranges_global_identity_replay_and_unchanged_status(setup):
    db, actor, service = setup
    request, result = add(db, actor, service)
    assert result == {'new': 3, 'unchanged': 0, 'stock_posted': False}
    assert [row.code for row in service.locations(db)] == ['A01', 'A02', 'A03']
    assert service.apply(db, actor, request) == result
    with pytest.raises(HTTPException) as error:
        service.apply(db, actor, request.model_copy(update={'rows': body([{'warehouse':'其他仓', 'bins':'A01'}]).rows}))
    assert error.value.status_code == 409
    db.rollback()
    from app.schemas.fabric_master import SaveMasterRequest
    single = SaveMasterRequest(factory_id='huakang-c', request_id=uuid5(request.request_id, 'fabric-location-1'), kind='LOCATION', code='B01', name='B01', status='ACTIVE', data={'warehouse':'布料仓'})
    with pytest.raises(HTTPException) as error:
        service.master.save(db, actor, single)
    assert error.value.status_code == 409
    db.rollback()
    row = service.locations(db)[0]
    service.master.save(db, actor, SaveMasterRequest(factory_id='huakang-c', request_id=uuid4(), id=row.id, expected_revision=row.revision, kind='LOCATION', code=row.code, name=row.name, status='INACTIVE', data={'warehouse':'布料仓'}))
    plan = service.preview(db, actor, body([{'warehouse':'布料仓','bins':'A01-A04'}]))
    assert plan['new'] == 1 and plan['unchanged'] == 3 and plan['items'][0]['status'] == 'INACTIVE'
    assert service.preview(db, actor, body([{'warehouse':'另一仓','bins':'a01'}]))['errors']
    with pytest.raises(HTTPException) as error:
        service.master.save(db, actor, single.model_copy(update={'request_id':uuid4(), 'code':'a01'}))
    assert error.value.status_code == 409
    db.rollback()
    for table in ('fabric_receipts','fabric_stock_batches','fabric_inventory_movements'):
        assert db.scalar(text(f'SELECT count(*) FROM {table}')) == 0


def test_stale_preview_and_conflict_do_not_partially_save(setup):
    db, actor, service = setup
    payload = body([{'warehouse':'布料仓','bins':'A01-A03'}])
    before = service.preview(db, actor, payload)
    add(db, actor, service, bins='B01')
    with pytest.raises(HTTPException) as error:
        service.apply(db, actor, apply_body(payload.rows, before['preview_token']))
    assert error.value.status_code == 409
    db.rollback()
    rows = body([{'warehouse':'布料仓','bins':'C01'}, {'warehouse':'其他仓','bins':'B01'}])
    plan = service.preview(db, actor, rows)
    assert plan['errors']
    with pytest.raises(HTTPException) as error:
        service.apply(db, actor, apply_body(rows.rows, plan['preview_token']))
    assert error.value.status_code == 422
    db.rollback()
    assert [row.code for row in service.locations(db)] == ['B01']


def test_warehouse_rename_checks_complete_group_preserves_codes_and_replays(setup):
    db, actor, service = setup
    add(db, actor, service)
    group = service.locations(db)
    payload = rename_body(group)
    add(db, actor, service, bins='A04')
    with pytest.raises(HTTPException) as error:
        service.rename(db, actor, payload)
    assert error.value.status_code == 409
    db.rollback()
    group = service.locations(db)
    before = {row.id: row.code for row in group}
    payload = rename_body(group)
    result = service.rename(db, actor, payload)
    assert result['updated'] == 4 and not result['stock_posted']
    assert {row.id: row.code for row in service.locations(db)} == before
    assert all(json.loads(row.data_json)['warehouse'] == '新布料仓' for row in service.locations(db))
    assert service.rename(db, actor, payload) == result
    assert len(service.master.history(db, group[0].id)['items']) == 2
    add(db, actor, service, warehouse='其他仓', bins='B01')
    conflict = payload.model_copy(update={'request_id':uuid4(), 'warehouse':'新布料仓', 'name':'其他仓', 'expected_group_token':service.master.warehouse_tokens(service.locations(db))['新布料仓']})
    with pytest.raises(HTTPException) as error:
        service.rename(db, actor, conflict)
    assert error.value.status_code == 409


def test_excel_template_ranges_formula_and_numeric_identifiers(setup):
    db, actor, service = setup
    book = load_workbook(BytesIO(service.template()))
    assert book['导入数据'].max_row == 1
    book['导入数据'].append(['布料仓','A01-A03'])
    output = BytesIO(); book.save(output)
    rows = service.file_rows(output.getvalue(), 'locations.xlsx')
    assert [row.bins for row in rows] == ['A01','A02','A03']
    assert service.preview(db, actor, body(rows))['new'] == 3
    for value in ('=1+1', 1):
        book['导入数据']['B2'] = value
        output = BytesIO(); book.save(output)
        with pytest.raises(HTTPException) as error:
            service.file_rows(output.getvalue(), 'locations.xlsx')
        assert error.value.status_code == 422
    assert service.preview(db, actor, body([{'warehouse':'布料仓','bins':'A1-A1001'}]))['errors']
    assert service.preview(db, actor, body([{'warehouse':'布料仓','bins':'A01、a01'}]))['errors']


def test_warehouse_with_more_than_one_batch_can_rename(setup):
    db, actor, service = setup
    add(db, actor, service, bins='A0001-A1000')
    add(db, actor, service, bins='B01')
    group = service.locations(db)
    assert len(group) == 1001
    assert service.rename(db, actor, rename_body(group))['updated'] == 1001


def test_visible_list_and_group_token_share_one_snapshot(setup, monkeypatch):
    db, actor, service = setup
    add(db, actor, service, bins='A01')
    original_records = service.master.records
    changed = False
    def interleaved_records(session):
        nonlocal changed
        rows = original_records(session)
        if session is db and not changed:
            changed = True
            with Session(db.get_bind()) as other:
                add(other, actor, service, bins='B01')
        return rows
    monkeypatch.setattr(service.master, 'records', interleaved_records)
    listed = service.master.list_records(db, 'LOCATION')
    assert [item['code'] for item in listed['items']] == ['A01']
    payload = rename_body(original_records(db))
    payload = payload.model_copy(update={'expected_group_token':listed['warehouse_tokens']['布料仓']})
    with pytest.raises(HTTPException) as error:
        service.rename(db, actor, payload)
    assert error.value.status_code == 409


def test_template_uppercase_warehouse_keeps_existing_manual_name(setup):
    db, actor, service = setup
    add(db, actor, service, warehouse='Fabric-a', bins='A01')
    plan = service.preview(db, actor, body([{'warehouse':'FABRIC-A','bins':'A01-A02'}]))
    assert not plan['errors'] and plan['new'] == 1 and plan['unchanged'] == 1
    assert all(row['warehouse'] == 'Fabric-a' for row in plan['items'])


def test_batch_write_failure_rolls_back_all_records_and_evidence(setup, monkeypatch):
    db, actor, service = setup
    payload = body([{'warehouse':'布料仓','bins':'A01-A03'}])
    plan = service.preview(db, actor, payload)
    original_flush = db.flush
    calls = 0
    def fail_second(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise RuntimeError('simulated write failure')
        return original_flush(*args, **kwargs)
    monkeypatch.setattr(db, 'flush', fail_second)
    with pytest.raises(RuntimeError):
        service.apply(db, actor, apply_body(payload.rows, plan['preview_token']))
    db.rollback()
    monkeypatch.setattr(db, 'flush', original_flush)
    assert not service.locations(db)
    assert db.scalar(text('SELECT count(*) FROM fabric_master_changes')) == 0


def test_location_api_auth_file_flow_and_receipt_snapshot(monkeypatch):
    from dataclasses import replace
    from test_molding_sample_api import make_client, login_as
    from test_fabric_procurement import BASE, book, row, stored
    from test_fabric_procurement_tracking import save
    from test_fabric_receiving import receive, request
    from test_fabric_master_chase import master
    root='/api/fabric-warehouse/master'
    with make_client(monkeypatch) as client:
        login_as(client,'admin')
        assert client.get(root+'/locations/template',params={'factory_id':'huakang-c'}).status_code == 200
        template = client.get(root+'/locations/template',params={'factory_id':'huakang-c'}).content
        imported = load_workbook(BytesIO(template))
        imported['导入数据'].append(['布料仓','C01-C03'])
        buffer=BytesIO(); imported.save(buffer); imported.close()
        file_plan=client.post(root+'/locations/file-preview',data={'factory_id':'huakang-c'},files={'file':('locations.xlsx',buffer.getvalue())})
        assert file_plan.status_code == 200,file_plan.text
        file_body={'factory_id':'huakang-c','rows':file_plan.json()['rows'],'preview_token':file_plan.json()['preview_token'],'request_id':str(uuid4())}
        assert client.post(root+'/locations/apply',json=file_body).json()['new'] == 3
        body={'factory_id':'huakang-c','rows':[{'warehouse':'布料仓','bins':'A01-A02'}]}
        plan=client.post(root+'/locations/preview',json=body)
        assert plan.status_code == 200,plan.text
        submitted={**body,'request_id':str(uuid4()),'preview_token':plan.json()['preview_token']}
        result=client.post(root+'/locations/apply',json=submitted)
        assert result.status_code == 200,result.text
        assert client.post(root+'/locations/apply',json=submitted).json() == result.json()
        save(client,book([row(采购明细ID='L1',入库数量=0,交货明细='')]))
        assert client.post(root,json=master()).status_code == 200
        line=stored(client)['items'][0]
        bins=client.get(root,params={'factory_id':'huakang-c','kind':'LOCATION'}).json()['items']
        location_id=next(r['id'] for r in bins if r['code'] == 'A01')
        posted=receive(client,line,request(line,batches=[{'quantity':'10','location':'a01','location_id':location_id,'dye_lot':'001'}]))
        assert posted.status_code == 200,posted.text
        assert posted.json()['batches'][0]['location'] == '布料仓／A01'
        locations=client.get(root,params={'factory_id':'huakang-c','kind':'LOCATION'}).json()['items']
        rename={'factory_id':'huakang-c','request_id':str(uuid4()),'warehouse':'布料仓','name':'布料新仓','expected_group_token':client.get(root,params={'factory_id':'huakang-c','kind':'LOCATION'}).json()['warehouse_tokens']['布料仓']}
        assert client.post(root+'/locations/rename-warehouse',json=rename).status_code == 200
        historical=client.get(BASE+'/receipts',params={'factory_id':'huakang-c'}).json()['items'][0]
        assert historical['master_references']['locations'][0]['data']['warehouse'] == '布料仓'
        latest=client.get(root,params={'factory_id':'huakang-c','kind':'LOCATION'}).json()['items']
        manage={'factory_id':'huakang-c','request_id':str(uuid4()),'action':'DELETE','selections':[{'id':r['id'],'expected_revision':r['revision']} for r in latest]}
        blocked=client.post(root+'/locations/manage',json=manage)
        assert blocked.status_code == 409 and 'A01' in blocked.text
        assert len(client.get(root,params={'factory_id':'huakang-c','kind':'LOCATION'}).json()['items']) == 5
        unused=next(r for r in latest if r['code'] == 'C01')
        delete_one={**manage,'request_id':str(uuid4()),'selections':[{'id':unused['id'],'expected_revision':unused['revision']}]}
        assert client.post(root+'/locations/manage',json=delete_one).json()['updated'] == 1
        assert len(client.get(root,params={'factory_id':'huakang-c','kind':'LOCATION'}).json()['deleted_items']) == 1
        current=client.get(root,params={'factory_id':'huakang-c','kind':'LOCATION'}).json()['items'][0]
        assert client.post(root,json=master('LOCATION',current['code'],current['name'],id=current['id'],expected_revision=current['revision'],status='INACTIVE',data=current['data'])).status_code == 200
        line=stored(client)['items'][0]
        assert receive(client,line,request(line,delivery_reference='DN2',batches=[{'quantity':'1','location':'A01','location_id':location_id,'dye_lot':'002'}])).status_code == 422
        assert receive(client,line,request(line,delivery_reference='DN3',batches=[{'quantity':'1','location':'a01','location_id':location_id,'dye_lot':'002'}])).status_code == 422
        auth=importlib.import_module('app.services.auth')
        dbm=importlib.import_module('app.db')
        users=importlib.import_module('app.models.auth')
        with dbm.SessionLocal() as db:
            actor=auth.build_auth_context(db,db.get(users.AuthUser,'user-admin'))
        denied=replace(actor,overrides=(auth.AuthOverrideContext(id='deny',permission_code='fabric_warehouse:import',effect='deny',factory_id='huakang-c',department='pmc-warehouse'),))
        client.app.dependency_overrides[auth.get_current_user]=lambda:denied
        assert client.post(root+'/locations/preview',json=body).status_code == 403
        assert client.post(root+'/locations/apply',json=submitted).status_code == 403
        assert client.post(root+'/locations/rename-warehouse',json=rename).status_code == 403
        assert client.post(root+'/locations/manage',json=delete_one).status_code == 403
        assert client.post(root+'/locations/manage',json={**delete_one,'factory_id':'huakang-d'}).status_code == 422
        assert client.post(root+'/locations/file-preview',data={'factory_id':'huakang-c'},files={'file':('locations.xlsx',b'bad')}).status_code == 403
        assert client.post(root+'/locations/preview',json={**body,'factory_id':'huakang-d'}).status_code == 422
