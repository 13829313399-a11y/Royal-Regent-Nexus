from io import BytesIO
from decimal import Decimal as D
from openpyxl import Workbook
from test_spray_production import client, post, collection, setup_order, plan


def test_diamond_graph_partial_transfer_rework_and_conservation(client):
    c=client
    nodes=[{"key":"A","name":"底色","capability":"manual","predecessors":[]},
           {"key":"B","name":"左眼","capability":"manual","predecessors":["A"]},
           {"key":"C","name":"右眼","capability":"manual","predecessors":["A"]},
           {"key":"D","name":"检查","capability":"manual","predecessors":["B","C"]}]
    post(c,"orders",{"document_no":"DAG-001","customer":"华康B","due_date":"2026-09-30","lines":[{"product_no":"0001-1","part_name":"头","quantity":"100","steps":nodes}]})
    line=collection(c,"lines")[0];steps=sorted(collection(c,"steps"),key=lambda x:x["sequence"])
    post(c,"prepare",{"line_id":line["id"],"ready_at":"2026-08-01T00:00:00Z","reason":"首件确认"})
    res=post(c,"resources",{"name":"喷台","capability":"manual","workers":1,"machine_rate":100,"person_minutes":".6","calendar":[{"start":"2026-08-01T00:00:00Z","end":"2026-12-01T00:00:00Z"}]})
    batch=post(c,"receipts",{"document_no":"IN-DAG","business_date":"2026-09-11","lines":[{"source_line":"1","line_id":line["id"],"accepted":100}]})
    def attempt(step,state,qty=100,good=100,rework=0):
        p=post(c,"schedule-preview",{"changes":[{"batch_id":batch["id"],"step_id":step["id"],"resource_id":res["id"],"quantity":qty,"workers":1,"start_at":"2026-09-11T00:00:00Z","input_state":state}]})
        t=post(c,"schedule-apply",{"scenario_id":p["id"]})["ids"][0]
        post(c,"task-state",{"task_id":t,"action":"start"})
        post(c,"reports",{"business_date":"2026-09-11","shift":"白班","team":"一组","lines":[{"task_id":t,"regular_qty":qty,"good":good,"rework":rework,"people":[{"name":"甲"}]}]})
    attempt(steps[0],"route:")
    route="route:"+steps[0]["id"]
    post(c,"schedule-preview",{"changes":[{"batch_id":batch["id"],"step_id":steps[3]["id"],"resource_id":res["id"],"quantity":100,"workers":1,"start_at":"2026-09-11T00:00:00Z","input_state":route}]},expected=409)
    attempt(steps[1],route,good=90,rework=10)
    defect="rework:"+steps[1]["id"]+":"+route
    attempt(steps[1],defect,qty=10,good=10)
    route="route:"+",".join(sorted([steps[0]["id"],steps[1]["id"]]))
    attempt(steps[2],route)
    route="route:"+",".join(sorted(x["id"] for x in steps[:3]))
    attempt(steps[3],route)
    balances=c.get(f'/api/spray-production/batches/{batch["id"]}/trace',params={"factory_id":"huaxing"}).json()["balances"]
    assert D(balances["finished"])==100
    assert sum(D(x) for x in balances.values())==100


def test_graph_cycle_rejected_atomically(client):
    post(client,"orders",{"document_no":"CYCLE","customer":"甲","due_date":"2026-09-30","lines":[{"product_no":"01","part_name":"头","quantity":1,"steps":[{"key":"a","name":"A","capability":"manual","predecessors":["b"]},{"key":"b","name":"B","capability":"manual","predecessors":["a"]}]}]},expected=422)
    assert not collection(client,"orders")


def test_cross_shift_split_and_shared_person_collision(client):
    c=client;line,steps,_,batch=setup_order(c)
    windows=[{"start":"2026-09-12T08:00:00+08:00","end":"2026-09-12T10:00:00+08:00"},{"start":"2026-09-13T08:00:00+08:00","end":"2026-09-13T12:00:00+08:00"}]
    person=post(c,"shared-resources",{"kind":"person","name":"甲","calendar":windows})
    for name in ["共享台 A","共享台 B"]:
        post(c,"resources",{"name":name,"capability":"manual","workers":1,"machine_rate":1000,"person_minutes":".06","calendar":windows,"shared_requirements":[{"id":person["id"],"quantity":1}]})
    a,b=[x for x in collection(c,"resources") if x["name"].startswith("共享台")]
    base={"batch_id":batch["id"],"step_id":steps[0]["id"],"quantity":1000,"workers":1,"start_at":windows[0]["start"]}
    post(c,"schedule-preview",{"changes":[{**base,"resource_id":a["id"]},{**base,"resource_id":b["id"]}]},expected=409)
    # Remove the unrelated broad calendar resource from the suggestion candidates in this test.
    from app.db import SessionLocal
    from app.models.spray_production import SprayResource
    from sqlalchemy import select
    with SessionLocal() as db:
        original=db.scalar(select(SprayResource).where(SprayResource.name=="手喷工位 01"));original.calendar=[];db.commit()
    suggestion=post(c,"schedule-suggest",{"start_at":windows[0]["start"]})
    assert len(suggestion["changes"])==2
    assert sum(D(v["quantity"]) for v in suggestion["changes"])==4900
    assert suggestion["changes"][0]["end_at"]<=suggestion["changes"][1]["start_at"]
    post(c,"schedule-apply",{"scenario_id":suggestion["id"]})
    tasks=sorted(collection(c,'tasks'),key=lambda x:x['start_at'])
    changes=[{**task,'task_id':task['id'],'expected_revision':task['revision'],'quantity':qty} for task,qty in zip(tasks,[1000,3900])]
    moved=post(c,'schedule-preview',{'changes':changes})
    post(c,'schedule-apply',{'scenario_id':moved['id']})
    assert sum(D(t['quantity']) for t in collection(c,'tasks') if t['status']=='planned')==4900


def test_employee_multiple_tasks_one_subsidy_and_frozen_payroll(client):
    c=client
    rate=post(c,"rates",{"name":"计时版本","rule_code":"time_based","quantity_basis":"good","currency":"CNY","effective_date":"2026-01-01","parameters":{"rate":10},"evidence":"测试正式规则"})
    policy=post(c,"rates",{"name":"员工班次保底","rule_code":"employee_shift_guarantee","quantity_basis":"good","currency":"CNY","effective_date":"2026-01-01","parameters":{"regular_hour_rate":15,"overtime_hour_rate":20},"evidence":"测试负责人确认每班汇总一次"})
    for r in [rate,policy]:post(c,"rate-activate",{"rate_id":r["id"],"reason":"测试授权版本"})
    for activity,hours in [("调机",3),("样板",2)]:
        post(c,"reports",{"business_date":"2026-09-11","shift":"白班","team":"一组","lines":[{"activity":activity,"regular_hours":hours,"people":[{"name":"员工甲"}],"rate_id":rate["id"]}]})
    payload={"factory_id":"huaxing","business_date":"2026-09-11","shift":"白班","policy_id":policy["id"]}
    preview=c.post('/api/spray-production/payroll/preview',json=payload).json()
    assert preview["confirmable"]
    e=preview["employees"][0]
    assert D(e["earned"])==50 and D(e["subsidy"])==25 and D(e["payable"])==75
    post(c,"payroll-confirm",{**payload,"fingerprint":preview["fingerprint"]},key="one-payroll")
    post(c,"payroll-confirm",{**payload,"fingerprint":preview["fingerprint"]},key="one-payroll")
    assert len(collection(c,"payrolls"))==1
    post(c,"reports",{"business_date":"2026-09-11","shift":"白班","team":"二组","lines":[{"activity":"调机","regular_hours":1,"people":[{"name":"员工甲"}]}]},expected=409)


def workbook(amount):
    wb=Workbook();ws=wb.active;ws.title="送货"
    ws.append(["日期","送货单","货号","部件","数量","金额"])
    ws.append(["2020-08-11","001013170","0001-1","头",100,amount])
    stream=BytesIO();wb.save(stream);return stream.getvalue()


def test_history_dedup_difference_reversal_and_no_live_stock(client):
    c=client
    def upload(amount,name):return c.post('/api/spray-production/imports',params={"factory_id":"huaxing"},files={"file":(name,workbook(amount))}).json()["id"]
    source=upload(24,"历史.xlsx")
    fields={"business_date":"A","document_no":"B","product_no":"C","part_name":"D","quantity":"E","amount":"F","customer":{"value":"华康B"},"unit":{"value":"件"},"currency":{"value":"HKD"}}
    def apply(ident,approve=False):
        p={"factory_id":"huaxing","import_id":ident,"sheets":[{"sheet":"送货","kind":"shipment","role":"source","start_row":2,"end_row":2,"fields":fields}]}
        preview=c.post('/api/spray-production/history/preview',json=p).json()
        selected=[f'{r["row_id"]}:{r["variant"]}' for r in preview["rows"]]
        return post(c,"history-apply",{**p,"fingerprint":preview["fingerprint"],"selected":selected,"accept_differences":selected if approve else [],"reason":"测试华兴历史归属"}),preview
    result,_=apply(source);assert result["counts"]["added"]==1
    assert collection(c,"history-facts")[0]["document_no"]=="001013170"
    assert not collection(c,"batches")
    result,_=apply(source);assert result["counts"]["linked"]==1
    changed=upload(30,"历史.xlsx")
    result,preview=apply(changed);assert preview["counts"]["different"]==1 and result["counts"]["pending"]==1
    result,_=apply(changed,True);assert result["counts"]["corrected"]==1
    assert D(collection(c,"history-facts")[0]["values"]["amount"])==30
    post(c,"history-reverse",{"import_id":changed,"reason":"撤回更正版"})
    assert D(collection(c,"history-facts")[0]["values"]["amount"])==24
    assert not collection(c,"batches","huakang-b")


def test_persistent_draft_revision_and_factory_isolation(client):
    c=client
    p={"kind":"report_draft","name":"白班","payload":{"lines":[{"activity":"调机","regular_hours":"3"}]},"expected_revision":0}
    post(c,"preference-save",p)
    stored=c.get('/api/spray-production/preferences',params={"factory_id":"huaxing"}).json()["items"]
    assert stored[0]["payload"]==p["payload"]
    assert not c.get('/api/spray-production/preferences',params={"factory_id":"huakang-b"}).json()["items"]
    post(c,"preference-save",p,expected=409)
    post(c,"preference-save",{**p,"expected_revision":stored[0]["revision"],"payload":{"lines":[]}})


def test_operating_cycle_material_payroll_delivery_return_and_next_period_adjustment(client):
    c=client;line,steps,res,batch=setup_order(c)
    purchase=post(c,"purchases",{"document_no":"PO-01","supplier":"测试油漆商","material":"测试油漆","quantity":2,"unit":"桶","base_unit":"KG","conversion":160,"conversion_evidence":"测试包装 160KG/桶","price":1280,"currency":"CNY"})
    for index,(kind,qty) in enumerate([("receipt",320),("issue",30),("consume",20),("return",10)]):
        post(c,"material-events",{"purchase_id":purchase["id"],"kind":kind,"quantity":qty,"business_date":"2026-09-11","document_no":"MAT-"+str(index),"order_id":line["order_id"]})
    assert sum(D(e["cost"]) for e in collection(c,"material-events"))==160
    rate=post(c,"rates",{"name":"测试计件","rule_code":"piece_direct","quantity_basis":"good","currency":"CNY","effective_date":"2026-01-01","parameters":{"rate":".1"},"evidence":"明确测试计件规则"})
    policy=post(c,"rates",{"name":"测试班次保底","rule_code":"employee_shift_guarantee","quantity_basis":"good","currency":"CNY","effective_date":"2026-01-01","parameters":{"regular_hour_rate":15,"overtime_hour_rate":20},"evidence":"明确测试班次口径"})
    for r in [rate,policy]:post(c,"rate-activate",{"rate_id":r["id"],"reason":"测试负责人确认"})
    task=plan(c,batch,steps[0],res,100)
    post(c,"reports",{"business_date":"2026-09-11","shift":"白班","team":"一组","lines":[{"task_id":task,"regular_qty":100,"regular_hours":1,"good":100,"people":[{"name":"甲"}],"rate_id":rate["id"]}]})
    p={"factory_id":"huaxing","business_date":"2026-09-11","shift":"白班","policy_id":policy["id"]}
    wage=c.post('/api/spray-production/payroll/preview',json=p).json()
    post(c,"payroll-confirm",{**p,"fingerprint":wage["fingerprint"]})
    assert D(collection(c,"payroll-lines")[0]["payable"])==15
    post(c,"shipments",{"document_no":"SHIP-001","customer":"华康B","business_date":"2026-09-11","lines":[{"batch_id":batch["id"],"quantity":100}]})
    preview=c.post('/api/spray-production/settlements/preview',json={"factory_id":"huaxing","customer":"华康B","period":"2026-09","currency":"HKD"}).json()
    frozen=post(c,"settlements",{**preview,"line_ids":sorted(x["id"] for x in preview["lines"])})
    post(c,"settlement-adjust",{"settlement_id":frozen["id"],"business_date":"2026-10-01","amount":"-2","document_no":"ADJUST-01","reason":"已确认的价差"})
    shipment=collection(c,"shipment-lines")[0]
    post(c,"returns",{"shipment_line_id":shipment["id"],"quantity":10,"business_date":"2026-10-01","document_no":"RET-01","reason":"次月退回待判"})
    october=c.post('/api/spray-production/settlements/preview',json={"factory_id":"huaxing","customer":"华康B","period":"2026-10","currency":"HKD"}).json()
    assert D(october["amount"])==D('-4.40') and len(october["adjustments"])==1
    post(c,"settlements",{**october,"line_ids":sorted(x["id"] for x in october["lines"])})
    assert D(next(x for x in collection(c,"settlements") if x["id"]==frozen["id"])["amount"])==24


def test_original_template_preserves_layout_and_uses_factory_checked_data(client):
    from openpyxl import load_workbook
    c=client;line,steps,res,batch=setup_order(c)
    wb=Workbook();sheet=wb.active;sheet.title="送货单";sheet.merge_cells('A1:F1');sheet['A1']='原格式送货单';sheet.column_dimensions['A'].width=23;sheet.print_area='A1:F20'
    stream=BytesIO();wb.save(stream)
    source=c.post('/api/spray-production/imports',params={"factory_id":"huaxing"},files={"file":("template.xlsx",stream.getvalue())}).json()["id"]
    result=c.post(f'/api/spray-production/imports/{source}/render',json={"factory_id":"huaxing","bindings":[{"kind":"order_line","record_id":line["id"],"field":"product_no","sheet":"送货单","cell":"A4"},{"kind":"order_line","record_id":line["id"],"field":"quantity","sheet":"送货单","cell":"E4"}]})
    assert result.status_code==200,result.text[:300]
    rendered=load_workbook(BytesIO(result.content));ws=rendered['送货单']
    assert ws['A4'].value=='000935373-1' and ws['E4'].value==50000
    assert 'A1:F1' in str(ws.merged_cells) and ws.column_dimensions['A'].width==23 and 'F$20' in ws.print_area
    assert c.get(f'/api/spray-production/imports/{source}/template',params={"factory_id":"huaxing"}).content==stream.getvalue()
    assert c.post(f'/api/spray-production/imports/{source}/render',json={"factory_id":"huakang-b","bindings":[]}).status_code==404


def test_template_refreshes_dependent_totals_and_preserves_identifier_and_precision():
    from types import SimpleNamespace
    from openpyxl import load_workbook
    from app.services.spray_templates import patch_workbook
    from app.services.spray_imports import parse_source
    from app.services.spray_history import cell_value
    wb=Workbook();ws=wb.active;ws['A1']=123;ws['A1'].number_format='000000';ws['B1']=2;ws['C1']='=B1*3';ws['D1']='=SUM(B1:C1)';ws['E1']='=UNSUPPORTED(B1)'
    buf=BytesIO();wb.save(buf)
    result=patch_workbook(SimpleNamespace(format='OOXML',content=buf.getvalue()),{('Sheet','B1'):('4',True)})
    cached=load_workbook(BytesIO(result),data_only=True).active
    assert cached['C1'].value==12 and cached['D1'].value==16 and cached['E1'].value is None
    _,rows=parse_source(result);cells=next(r['cells'] for r in rows if r['row_number']==1)
    row=SimpleNamespace(cells=cells,row_number=1)
    assert cell_value(row,'A','product_no')[0]=='000123'
    row.cells['F1']={'raw_value':'74.54545454545455'}
    assert cell_value(row,'F','wage')[0]=='74.54545454545455'
