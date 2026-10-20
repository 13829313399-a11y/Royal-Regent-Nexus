"""Synthetic Shushupapa source-to-preview-to-export integration contracts."""
from copy import deepcopy
from datetime import datetime
from io import BytesIO

import openpyxl
import pytest

from app.schemas.customer_order import CustomerOrderLineOut
from app.services import customer_order_unified as service
from app.services.huaxing_order_legacy import multi_po_parser, multi_schedule
from test_customer_order_shushupapa_export import amount_formula, find_row, workbook


ITEM = '87654321'
CONTRACT = '2098-0701'
CUSTOMER_PO = '801-9999-121'


def source_order():
    def evidence(page, raw_text):
        return dict(source_file='synthetic-source.pdf', page=page, raw_text=raw_text)

    return dict(
        client='shushupapa', client_name='SHUSHUPAPA',
        filename='synthetic-source.pdf', po_number=CONTRACT,
        customer_po=CUSTOMER_PO, customer='Costco', country='Japan',
        po_date='2028-06-09', ship_date='2028-07-21', source_currency='USD',
        warnings=['Synthetic document warning: confirm carton artwork.'],
        lines=[dict(
            item_code=ITEM, description='18 inch Example Figure, Orbit Pilot',
            quantity=180, outer_pack=12, cartons=15, unit_price=3.25, amount=585,
            source_currency='USD', source_page=1,
            standard='Example Standard X-42', color_box='Costco POD; 12 pcs/POD',
            box_mark='TBA', label='TBA',
            source_warnings=['Synthetic amount-in-words discrepancy; verify source.'],
            source_lineage={
                'po_date': evidence(1, 'Issue Date: Jun 9, 2028'),
                'customer_po': evidence(1, f'Costco PO Number: {CUSTOMER_PO}'),
                'unit_price': evidence(1, 'USD 3.25'),
                'amount': evidence(1, 'USD 585.00'),
                'outer_pack': evidence(2, 'Master Carton Packed with 12 pcs'),
                'ship_date': evidence(2, 'Delivery Date: Jul 21, 2028'),
                'color_box': evidence(1, 'Costco POD; 12 pcs/POD'),
                'standard': evidence(2, 'Example Standard X-42'),
            },
        )],
    )


def template_bytes(history_customer='Costco'):
    book = workbook()
    item = book['ITEM表']
    item['G4'], item['H4'], item['I4'] = history_customer, ITEM, '合成玩具历史中文名'
    item['N4'], item['T4'], item['U4'] = 'OLD STANDARD', 'OLD RETAILER PACKAGING', 'OLD-DATE-CODE'
    for name in ('接单表', '正单评审表'):
        sheet = book[name]
        sheet['F4'], sheet['G4'], sheet['N4'] = history_customer, ITEM, 99
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def arguments(history_customer='Costco'):
    return dict(
        customer_code='shushupapa', factory_id='huaxing', received_date='2028-06-10',
        po_files=[('synthetic-source.pdf', b'%PDF-1.4 synthetic test input')],
        schedule_file_name='synthetic-schedule.xlsx',
        schedule_content=template_bytes(history_customer),
    )


def mock_source(monkeypatch, order=None):
    original = source_order() if order is None else order
    monkeypatch.setattr(multi_po_parser, 'parse_po', lambda *args, **kwargs: deepcopy(original))
    return original


def override(field, value):
    return dict(row_id='shushupapa-1', field=field, value=value, reason='Synthetic manual confirmation')


def test_multi_schedule_carries_source_metadata_into_unified_preview_and_api(monkeypatch):
    order = mock_source(monkeypatch)
    record = multi_schedule._line_values(order, order['lines'][0])
    assert record['source_currency'] == 'USD'
    assert record['source_lineage'] == order['lines'][0]['source_lineage']
    assert record['source_warnings'] == order['lines'][0]['source_warnings']

    preview = service.create_unified_customer_preview(**arguments())
    row, = preview['rows']
    assert (row['contract_no'], row['reference_no'], row['po_no']) == (CONTRACT, CONTRACT, CUSTOMER_PO)
    assert (row['customer_name'], row['country'], row['standard']) == ('Costco', 'Japan', 'Example Standard X-42')
    assert (row['source_order_date'], row['requested_ship_date'], row['customer_q']) == (
        '2028-06-09', '2028-07-21', '2028-07-14',
    )
    assert (row['quantity'], row['units_per_carton'], row['carton_count']) == ('180', '12', '15')
    assert row['product_name_zh'] == '合成玩具历史中文名'
    assert row['packaging'] == 'Costco POD; 12 pcs/POD'
    assert row.get('date_code', '') == ''
    assert (row['unit_price_usd'], row['amount_usd']) == ('3.25', '585')
    assert row['unit_price_hkd'] == row['amount_hkd'] == ''
    assert '第 2 页' in row['lineage']['units_per_carton']
    assert 'Master Carton Packed with 12 pcs' in row['lineage']['units_per_carton']
    assert 'USD 3.25' in row['lineage']['unit_price_usd']
    assert any(issue['message'] == order['lines'][0]['source_warnings'][0] for issue in row['issues'])
    assert any(order['warnings'][0] in warning for warning in preview['warnings'])

    public_row = CustomerOrderLineOut.model_validate(row).model_dump(mode='json')
    assert public_row['unit_price_usd'] == '3.25'
    assert public_row['amount_usd'] == '585'
    assert public_row['unit_price_hkd'] == public_row['amount_hkd'] == ''
    assert public_row['lineage']['unit_price_usd'] == row['lineage']['unit_price_usd']


def test_new_retailer_does_not_inherit_other_retailer_pack_or_standard(monkeypatch):
    order = source_order()
    for field in ('outer_pack', 'cartons', 'standard', 'color_box'):
        order['lines'][0][field] = None
    mock_source(monkeypatch, order)
    row, = service.create_unified_customer_preview(**arguments('OTHER SYNTHETIC RETAILER'))['rows']
    assert row['customer_name'] == 'Costco'
    assert row['packaging'] == row['standard'] == row['units_per_carton'] == ''
    assert row['carton_count'] == ''
    assert row['unit_price_hkd'] == row['amount_hkd'] == ''
    assert row['product_name_zh'] == '合成玩具历史中文名'


@pytest.mark.parametrize('currency,price,expected', [('USD', 0, '0'), ('EUR', 3.25, '')])
def test_only_confirmed_usd_is_published_and_zero_is_preserved(monkeypatch, currency, price, expected):
    order = source_order()
    order['source_currency'] = order['lines'][0]['source_currency'] = currency
    order['lines'][0]['unit_price'] = price
    mock_source(monkeypatch, order)
    row, = service.create_unified_customer_preview(**arguments())['rows']
    assert row['unit_price_usd'] == expected
    assert row['amount_usd'] == ('0' if expected == '0' else '')
    assert row['unit_price_hkd'] == row['amount_hkd'] == ''


@pytest.mark.parametrize('original_ship', ['2028-07-21', None])
def test_manual_shipping_date_recomputes_planned_inspection_in_preview_and_workbook(monkeypatch, original_ship):
    order = source_order()
    order['ship_date'] = original_ship
    mock_source(monkeypatch, order)
    output, _, preview = service.export_unified_customer_schedule(
        **arguments(), manual_overrides=[override('requested_ship_date', '2028-08-10')],
    )
    row, = preview['rows']
    assert row['requested_ship_date'] == '2028-08-10'
    assert row['customer_q'] == '2028-08-03'
    assert '前 7 天' in row['lineage']['customer_q']
    sheet = openpyxl.load_workbook(BytesIO(output))['ITEM表']
    number = find_row(sheet, CONTRACT)
    assert sheet[f'Y{number}'].value == datetime(2028, 8, 3)
    assert sheet[f'Z{number}'].value == datetime(2028, 8, 10)


@pytest.mark.parametrize('manual_inspection', ['2028-07-14', '2028-08-01'])
def test_explicit_manual_inspection_is_retained_even_when_equal_to_original_derived_date(monkeypatch, manual_inspection):
    mock_source(monkeypatch)
    output, _, preview = service.export_unified_customer_schedule(
        **arguments(), manual_overrides=[override('requested_ship_date', '2028-08-10'),
                                        override('customer_q', manual_inspection)],
    )
    row, = preview['rows']
    assert row['customer_q'] == manual_inspection
    assert row['lineage']['customer_q'] == '业务人工补录计划验货日期'
    sheet = openpyxl.load_workbook(BytesIO(output))['ITEM表']
    assert sheet[f'Y{find_row(sheet, CONTRACT)}'].value == datetime.fromisoformat(manual_inspection)


def test_manual_quantity_recomputes_usd_and_export_links_both_summary_prices(monkeypatch):
    mock_source(monkeypatch)
    output, _, preview = service.export_unified_customer_schedule(
        **arguments(), manual_overrides=[override('quantity', '240'), override('unit_price_hkd', '25.5')],
    )
    row, = preview['rows']
    assert (row['quantity'], row['carton_count'], row['amount_usd'], row['amount_hkd']) == ('240', '20', '780', '6120')
    assert '金额 USD 780' in row['price_details_text']
    book = openpyxl.load_workbook(BytesIO(output))
    a, b = (find_row(book[name], CONTRACT) for name in ('接单表', '正单评审表'))
    assert book['接单表'][f'N{a}'].value == 25.5
    assert book['接单表'][f'P{a}'].value == 3.25
    for name, number in (('接单表', a), ('正单评审表', b)):
        assert book[name][f'O{number}'].value == amount_formula(number, 'N')
        assert book[name][f'Q{number}'].value == amount_formula(number, 'P')
    assert book['正单评审表'][f'N{b}'].value == f'=IF(LEN(\'接单表\'!N{a})=0,"",\'接单表\'!N{a})'
    assert book['正单评审表'][f'P{b}'].value == f'=IF(LEN(\'接单表\'!P{a})=0,"",\'接单表\'!P{a})'


def test_export_does_not_convert_usd_using_template_exchange_rate_or_history_price(monkeypatch):
    mock_source(monkeypatch)
    output, _, preview = service.export_unified_customer_schedule(**arguments())
    row, = preview['rows']
    assert row['unit_price_hkd'] == row['amount_hkd'] == ''
    book = openpyxl.load_workbook(BytesIO(output))
    a = find_row(book['接单表'], CONTRACT)
    assert book['接单表'][f'N{a}'].value is None
    assert book['接单表'][f'P{a}'].value == 3.25
    assert book['接单表']['N4'].value == 99


def test_usd_rounding_matches_excel_and_numeric_source_conflict_is_visible(monkeypatch):
    order = source_order()
    order['lines'][0].update(quantity=1, unit_price=2.345, amount=2.34)
    mock_source(monkeypatch, order)
    row, = service.create_unified_customer_preview(**arguments())['rows']
    assert row['amount_usd'] == '2.35'
    assert any(issue['code'] == 'shushupapa_amount_mismatch' for issue in row['issues'])
    assert '2.34' in row['lineage']['amount_usd']


@pytest.mark.parametrize('hkd_price,expected', [
    ('2.345', '2.35'), ('2.355', '2.36'), ('0', '0'), (None, ''),
])
def test_confirmed_hkd_rounding_matches_excel_without_inventing_missing_price(monkeypatch, hkd_price, expected):
    order = source_order()
    order['lines'][0].update(quantity=1, outer_pack=1, cartons=1, amount=3.25)
    mock_source(monkeypatch, order)
    overrides = [override('unit_price_hkd', hkd_price)] if hkd_price is not None else []

    output, _, preview = service.export_unified_customer_schedule(
        **arguments(), manual_overrides=overrides,
    )

    row, = preview['rows']
    assert row['amount_hkd'] == expected
    assert row['unit_price_hkd'] == (hkd_price if hkd_price is not None else '')
    assert row['unit_price_usd'] == row['amount_usd'] == '3.25'
    assert CustomerOrderLineOut.model_validate(row).amount_hkd == expected
    book = openpyxl.load_workbook(BytesIO(output))
    for name in ('接单表', '正单评审表'):
        number = find_row(book[name], CONTRACT)
        assert book[name][f'O{number}'].value == amount_formula(number, 'N')
    order_number = find_row(book['接单表'], CONTRACT)
    assert book['接单表'][f'N{order_number}'].value == (
        float(hkd_price) if hkd_price is not None else None
    )
