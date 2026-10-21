from copy import copy
from datetime import date
from io import BytesIO
import json
import os
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill, Font
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, TwoCellAnchor, AnchorMarker
from openpyxl.chart import BarChart, Reference
import pytest

from app.services import customer_order_unified as unified
from app.services.huaxing_order_legacy import edu_grouped_export as exporter


GREEN = PatternFill('solid', fgColor='FF4EA72E')
PRODUCT = 'F12-ZZ901-C901'


def template():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for index, name in enumerate(exporter.SHEETS):
        sheet = wb.create_sheet(name)
        for col, header in enumerate(unified.ITEM_HEADERS if name == 'ITEM表' else unified.ORDER_HEADERS, 1):
            sheet.cell(3, col, header)
        ref, item, quantity, po, contract = exporter._columns(sheet)
        sheet.cell(4, ref, 'EDUHX00010-1')
        sheet.cell(4, item, PRODUCT + '0001')
        sheet.cell(4, quantity, 20)
        sheet.cell(4, contract, 'OLD-CONTRACT')
        sheet.cell(4, po, 'OLD-PO')
        for cell in sheet[4]:
            cell.font = Font(name='Arial', size=10)
        sheet.cell(20 + index, 1, '取消单')
        sheet.cell(22 + index, 1, '已走货订单')
        sheet.cell(24 + index, item, 'OTHER-HISTORICAL')
        sheet.cell(24 + index, 31, '不能覆盖的人工扩展')
        sheet.print_area = f'A1:AE{30 + index}'
        sheet.freeze_panes = 'A4'
    wb['ITEM表']['L4'] = '1/16'
    wb['ITEM表']['M4'] = '=K4/16'
    wb['接单表']['P1'] = 7.6
    wb.create_sheet('已出发票')['XFD131'] = '原始尾列，不清理'
    return wb


def reserve(wb, r, *, product=PRODUCT + '0000', quantity='=3000-300', reference='EDUHX00010', po='RESERVE-PO'):
    sheet = wb['ITEM表']
    for col, value in {4:'PO订单', 5:reference, 6:po, 8:product, 11:quantity, 12:'1/16'}.items():
        sheet.cell(r, col, value)
    for cell in sheet[r]:
        cell.fill = copy(GREEN)


def order(**kw):
    result = dict(id='new', _huaxing_customer='edu', item_sheet_name='ITEM表',
                  product_no=PRODUCT, source_product_no=PRODUCT + '0003',
                  reference_no='EDUHX00104', contract_no='NEW-CONTRACT', po_no='NEW-PO',
                  quantity='300', units_per_carton='16', received_date='2026-10-16',
                  source_order_date='2026-10-09', product_name_zh='产品名', date_code='LOT 123', issues=[])
    result.update(kw)
    return result


def find_row(sheet, reference):
    ref = exporter._columns(sheet)[0]
    return next(r for r in range(4, sheet.max_row + 1) if sheet.cell(r, ref).value == reference)


def contents(wb):
    return sorted((s.title, c.coordinate, c.value, str(c._style)) for s in wb
                  for c in s._cells.values() if c.value is not None or c.has_style)


def test_grouped_insert_before_green_keeps_formula_and_historical_values():
    wb = template()
    reserve(wb, 6)
    old_note = wb['ITEM表']['AE24']
    old_formula = wb['ITEM表']['M4']
    report = exporter.export_rows(wb, [order(edu_reservation_row='6')])
    assert report['written'] == 1
    assert wb['ITEM表']['E6'].value == 'EDUHX00104'
    assert wb['ITEM表']['K7'].value == '=3000-300-300'
    assert not exporter.is_reservation_row(wb['ITEM表'], 6)
    assert exporter.is_reservation_row(wb['ITEM表'], 7)
    assert wb['ITEM表']['H6'].value == PRODUCT
    assert wb['ITEM表']['H6'].comment.text.endswith(PRODUCT + '0003')
    assert wb['ITEM表']['M6'].value == '=K6/L6'
    assert old_formula.value == '=K4/16'
    assert old_note.value == '不能覆盖的人工扩展' and old_note.row == 25
    assert wb['已出发票']['XFD131'].value == '原始尾列，不清理'
    for name in exporter.SHEETS[1:]:
        r = find_row(wb[name], 'EDUHX00104')
        assert r == 5
        assert wb[name].cell(r, 10).value == '=IF(LEN(\'ITEM表\'!K6)=0,"",\'ITEM表\'!K6)'


def test_selected_original_rows_survive_multiple_insertions_and_continuous_subtraction():
    wb = template()
    reserve(wb, 6, quantity=0)
    reserve(wb, 8, reference='EDUHX00011', quantity='=100-50')
    second = 'F12-SS909-C909'
    reserve(wb, 10, product=second + '0000', quantity=-100, reference='EDUHX00012')
    rows = [order(edu_reservation_row='8'),
            order(id='two', reference_no='EDUHX00105', contract_no='C2', quantity='400', edu_reservation_row='8'),
            order(id='three', reference_no='EDUHX00106', contract_no='C3', product_no=second,
                  source_product_no=second + '0000', quantity='500', edu_reservation_row='10')]
    exporter.export_rows(wb, rows)
    assert wb['ITEM表'].cell(find_row(wb['ITEM表'], 'EDUHX00010'), 11).value == 0
    assert wb['ITEM表'].cell(find_row(wb['ITEM表'], 'EDUHX00011'), 11).value == '=100-50-300-400'
    assert wb['ITEM表'].cell(find_row(wb['ITEM表'], 'EDUHX00012'), 11).value == '=-100-500'
    for name in exporter.SHEETS:
        assert find_row(wb[name], 'EDUHX00105') == find_row(wb[name], 'EDUHX00104') + 1


def test_duplicate_is_idempotent_even_without_preview_flag_and_conflict_is_atomic():
    wb = template()
    reserve(wb, 6)
    row = order(edu_reservation_row='6')
    exporter.export_rows(wb, [row])
    before = contents(wb)
    assert exporter.export_rows(wb, [row]) == {'written':0, 'skipped_existing':1, 'reservation_deductions':[]}
    assert before == contents(wb)
    with pytest.raises(unified.CustomerOrderUnifiedError, match='数量冲突'):
        exporter.export_rows(wb, [dict(row, quantity='301')])
    assert before == contents(wb)


def test_preview_existing_flag_skips_no_po_and_no_stock_change():
    wb = template()
    reserve(wb, 6)
    exporter.export_rows(wb, [order(_edu_existing_order=True, edu_reservation_row='6')])
    assert 'PO订单' not in wb.sheetnames
    assert wb['ITEM表']['K6'].value == '=3000-300'


@pytest.mark.parametrize('selected', ['', '5', '999', '6.0'])
def test_bad_reservation_selection_never_guesses_or_mutates(selected):
    wb = template()
    reserve(wb, 6)
    reserve(wb, 8, reference='EDUHX00011')
    before = contents(wb)
    with pytest.raises(unified.CustomerOrderUnifiedError, match='备料'):
        exporter.export_rows(wb, [order(edu_reservation_row=selected)])
    assert before == contents(wb)


def test_no_reservation_or_explicit_skip_adds_order_without_creating_stock():
    wb = template()
    reserve(wb, 6)
    exporter.export_rows(wb, [order(_edu_skip_reservation=True)])
    assert wb['ITEM表']['K7'].value == '=3000-300'
    exporter.export_rows(wb, [order(product_no='UNSEEN', reference_no='EDUHX00105', contract_no='C2')])
    assert all(c['product'] != 'UNSEEN' for c in exporter.reservation_candidates(wb, 'UNSEEN'))


def test_green_row_is_detected_without_po_label_but_pale_formula_color_is_not_stock():
    wb = template()
    reserve(wb, 6)
    wb['ITEM表']['D6'] = None
    assert exporter.reservation_candidates(wb, PRODUCT + '0008')[0]['row'] == 6
    sheet = wb['接单表']
    sheet['J4'] = '=ITEM表!K4'
    sheet['J4'].fill = PatternFill('solid', fgColor='FFEFF8EE')
    assert not exporter.is_reservation_row(sheet, 4)


def test_linked_summary_stock_is_not_double_deducted_literal_summary_budget_is_updated():
    wb = template()
    reserve(wb, 6, quantity=0)
    for name in exporter.SHEETS[1:]:
        sheet = wb[name]
        sheet['D7'], sheet['G7'] = 'EDUHX00010', PRODUCT + '0000'
        sheet['J7'] = '=IF(LEN(\'ITEM表\'!K6)=0,"",\'ITEM表\'!K6)' if name == '接单表' else 0
        for cell in sheet[7]: cell.fill = copy(GREEN)
    exporter.export_rows(wb, [order(edu_reservation_row='6')])
    assert wb['ITEM表']['K7'].value == '=0-300'
    assert wb['接单表']['J8'].value == '=IF(LEN(\'ITEM表\'!K7)=0,"",\'ITEM表\'!K7)'
    assert wb['正单评审表']['J8'].value == '=0-300'


def test_po_appends_source_date_normalized_product_and_blank_price_preserving_side_notes():
    wb = template()
    po = wb.create_sheet('PO订单')
    for c,h in enumerate(exporter.PO_HEADERS,1): po.cell(1,c,h)
    po.append([date(2026,1,1),'EDUHX00001',PRODUCT+'0000','原产品',100,9.99,'=E2*F2'])
    po['N3'] = '人工价格备注'; po['F2'].comment = Comment('价格不能覆盖','业务')
    exporter.export_rows(wb, [order(unit_price_hkd='12.5')])
    assert po['F2'].value == 9.99 and po['F2'].comment.text == '价格不能覆盖'
    assert po['N3'].value == '人工价格备注'
    assert po['A4'].value == date(2026,10,9)
    assert po['B4'].value == 'EDUHX00104' and po['C4'].value == PRODUCT
    assert po['F4'].value is None
    assert po['G4'].value == '=IF(OR(E4="",F4=""),"",E4*F4)'
    assert po['C4'].comment.text.endswith(PRODUCT+'0003')


def test_insertion_shifts_formulas_local_global_names_validation_cf_and_print_range():
    wb = template()
    reserve(wb, 6)
    item = wb['ITEM表']
    item['AE10'] = '=K6+K4'
    wb['已出发票']['A1'] = "='ITEM表'!K6"
    item.defined_names.add(DefinedName('LocalReservation', attr_text='$K$6'))
    wb.defined_names.add(DefinedName('AllReservation', attr_text="'ITEM表'!$K$6"))
    validation=DataValidation(type='custom', formula1='K6>=0'); validation.add('K6:K10')
    item.add_data_validation(validation)
    rule=FormulaRule(formula=['K6<0']); item.conditional_formatting.add('K6:K10',rule)
    item.freeze_panes='A7'
    exporter.export_rows(wb,[order(edu_reservation_row='6')])
    assert item['AE11'].value == '=K7+K4'
    assert wb['已出发票']['A1'].value == "='ITEM表'!K7"
    assert item.defined_names['LocalReservation'].attr_text == '$K$7'
    assert wb.defined_names['AllReservation'].attr_text == "'ITEM表'!$K$7"
    assert str(validation.sqref) == 'K7:K11' and validation.formula1=='K7>=0'
    assert str(next(iter(item.conditional_formatting)).sqref)=='K7:K11'
    assert rule.formula==['K7<0']
    assert item.freeze_panes=='A8' and '$AE$31' in str(item.print_area)


def test_batch_duplicates_only_create_one_po_row_and_one_deduction():
    wb=template(); reserve(wb,6)
    first=order(edu_reservation_row='6')
    report=exporter.export_rows(wb,[first,dict(first,id='same',reference_no='EDUHX00999')])
    assert report['written']==1 and report['skipped_existing']==1
    assert wb['ITEM表']['K7'].value=='=3000-300-300'
    assert wb['PO订单'].max_row==2


def test_existing_drawings_move_once_with_two_cell_size_preserved():
    wb=template(); reserve(wb,6)
    sheet=wb['ITEM表']
    one=BarChart(); one.anchor=OneCellAnchor(_from=AnchorMarker(col=1,row=9))
    two=BarChart(); two.anchor=TwoCellAnchor(_from=AnchorMarker(col=2,row=9),to=AnchorMarker(col=5,row=14))
    text=BarChart(); text.anchor='A10'
    for drawing in (one,two,text): sheet.add_chart(drawing)
    exporter.export_rows(wb,[order(edu_reservation_row='6')])
    assert one.anchor._from.row==10
    assert two.anchor._from.row==10 and two.anchor.to.row==15
    assert text.anchor=='A11'


def test_chart_source_references_follow_existing_data_even_on_another_sheet():
    wb=template(); reserve(wb,6)
    item=wb['ITEM表']
    for number in (7,8): item.cell(number,11,number*10)
    chart=BarChart()
    chart.add_data(Reference(item,min_col=11,min_row=6,max_row=8))
    chart.set_categories(Reference(item,min_col=8,min_row=6,max_row=8))
    wb['已出发票'].add_chart(chart,'A6')
    exporter.export_rows(wb,[order(edu_reservation_row='6')])
    assert chart.series[0].val.numRef.f=="'ITEM表'!$K$7:$K$9"
    assert chart.series[0].cat.numRef.f=="'ITEM表'!$H$7:$H$9"
    assert chart.anchor=='A6'


def test_private_schedule_preserves_all_existing_literals_styles_and_invoice_sparse_tail(tmp_path):
    configured = os.environ.get("HUAXING_EDU_GROUPED_TEST_MANIFEST")
    if not configured:
        pytest.skip("Set HUAXING_EDU_GROUPED_TEST_MANIFEST to a private local JSON manifest")
    manifest_path = Path(configured)
    case = json.loads(manifest_path.read_text(encoding="utf-8"))
    source = Path(case["schedule_path"])
    if not source.is_absolute():
        source = manifest_path.parent / source
    original = source.read_bytes()
    wb = openpyxl.load_workbook(BytesIO(original))
    from openpyxl.cell.cell import MergedCell
    old = [(c, c.value, copy(c._style)) for s in wb for c in s._cells.values() if not isinstance(c, MergedCell)]
    # Private fixtures define both explicit reservation choices and expected rows.
    specs = case["orders"]
    assert specs and all({"reference_no", "product_no", "contract_no", "quantity"} <= spec.keys() for spec in specs)
    rows = [order(id=str(i), **spec) for i, spec in enumerate(specs)]
    changed = {id(wb['ITEM表'].cell(int(row["edu_reservation_row"]), 11))
               for row in rows if row.get("edu_reservation_row")}
    for check in case["reservation_checks"]:
        assert len(exporter.reservation_candidates(wb, check["product_no"])) == check["expected_count"]
    report = exporter.export_rows(wb, rows)
    assert report['written'] == len(rows)
    for cell, value, style in old:
        assert cell._style == style
        if cell.data_type != 'f' and id(cell) not in changed:
            assert cell.value == value
    sparse = case["sparse_tail"]
    assert wb[sparse["sheet"]].max_column == sparse["expected_max_column"]
    for row in rows:
        ir = find_row(wb['ITEM表'], row['reference_no'])
        assert not exporter.is_reservation_row(wb['ITEM表'], ir)
        for sn in exporter.SHEETS[1:]:
            sr = find_row(wb[sn], row['reference_no'])
            assert f"'ITEM表'!K{ir}" in wb[sn].cell(sr, 10).value
    result = tmp_path / 'grouped-roundtrip.xlsx'
    wb.save(result)
    wb.close()
    reopened = openpyxl.load_workbook(result)
    try:
        assert exporter.export_rows(reopened, rows)['written'] == 0
    finally:
        reopened.close()
    assert source.read_bytes() == original
