"""Synthetic, independently asserted PDF content and visible geometry truth."""
import io

import pdfplumber
import pypdfium2 as pdfium
import pytest
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, FloatObject, NameObject, RectangleObject

from app.services.document_tools.document_ir import Cancelled, ToolError
from app.services.document_tools.pdf_engine import convert_pdf, inspect_pdf, split_suggestions
from app.services.document_tools.pdf_geometry import geometry, normalize_pdf, parse_page_groups, transform_rect


def make_pdf(path, *, rotation=0, unit=1, crop=None, pages=1, width=400, height=600, table=False, borderless=False):
    writer=PdfWriter()
    for n in range(pages):
        page=writer.add_blank_page(width,height)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        commands=[f'BT /F1 12 Tf 50 {height-60} Td (PAGE-{n+1:02d} 00123) Tj ET']
        if table or borderless:
            for top in [height-120,height-320]:
                if not borderless:
                    for x in [40,140,240,340]:
                        commands.append(f'{x} {top-90} m {x} {top} l S')
                    for y in [top,top-30,top-60,top-90]:
                        commands.append(f'40 {y} m 340 {y} l S')
                for ri,row in enumerate([['ITEM','QTY','AMOUNT'],['00123','12','12.50'],['12345678901234567','24','1,234']]):
                    for ci,text in enumerate(row):
                        commands.append(f'BT /F1 8 Tf {45+100*ci} {top-20-30*ri} Td ({text}) Tj ET')
        stream=DecodedStreamObject(); stream.set_data('\n'.join(commands).encode())
        page[NameObject('/Contents')]=writer._add_object(stream)
        page.rotate(rotation)
        page[NameObject('/UserUnit')]=FloatObject(unit)
        if crop:
            page.cropbox=RectangleObject(crop)
    with path.open('wb') as f:
        writer.write(f)
    return path


@pytest.mark.parametrize('rotation',[0,90,180,270])
@pytest.mark.parametrize('unit,crop',[(1,None),(1,[20,30,380,570]),(2,[20,30,380,570]),(.75,[20,30,380,570])])
def test_visible_geometry_and_cut_coverage(tmp_path,rotation,unit,crop):
    source=make_pdf(tmp_path/'source.pdf',rotation=rotation,unit=unit,crop=crop)
    g=geometry(PdfReader(source).pages[0])
    mapped=transform_rect(g['effective_crop_box'],g['raw_to_visible'])
    assert mapped==pytest.approx([0,0,g['width_pt'],g['height_pt']])
    cut=g['height_pt']*.413
    result=convert_pdf(source,'pdf_split',{'split_mode':'crop','axis':'y','cuts_pt':[cut]},tmp_path/'result',lambda *_:None,lambda:False)
    output=PdfReader(result.files[0]['path'])
    assert len(output.pages)==2
    assert float(output.pages[0].mediabox.height)==pytest.approx(cut,abs=.001)
    assert float(output.pages[1].mediabox.height)==pytest.approx(g['height_pt']-cut,abs=.001)
    assert sum(float(p.mediabox.height) for p in output.pages)==pytest.approx(g['height_pt'])
    assert output.pages[0].extract_text().count('00123')==1
    assert '/Font' in output.pages[0]['/Resources']
    with pdfium.PdfDocument(str(result.files[0]['path'])) as rendered:
        for i,page in enumerate(rendered):
            assert page.get_size()[1]==pytest.approx(float(output.pages[i].mediabox.height),abs=.001)
            bitmap=page.render(scale=.5)
            assert bitmap.width>0 and bitmap.height>0
            bitmap.close(); page.close()


def test_pdf_native_two_tables_editable_outputs(tmp_path):
    source=make_pdf(tmp_path/'tables.pdf',table=True)
    result=inspect_pdf(source,{},tmp_path/'inspect',lambda *_:None,lambda:False)
    assert len(result.ir.tables)==2
    assert [(t.row_count,t.column_count) for t in result.ir.tables]==[(3,3),(3,3)]
    assert result.ir.tables[0].cells[3].raw_text=='00123'
    assert all(c.source.bbox_pt and c.source.anchor_precision=='cell' for t in result.ir.tables for c in t.cells)
    word=convert_pdf(source,'pdf_to_word',{'ai_mode':'off'},tmp_path/'word',lambda *_:None,lambda:False)
    from docx import Document
    doc=Document(word.files[0]['path'])
    assert len(doc.tables)==2
    assert doc.tables[0].cell(1,0).text=='00123'
    excel=convert_pdf(source,'pdf_to_excel',{'ai_mode':'off'},tmp_path/'excel',lambda *_:None,lambda:False)
    from openpyxl import load_workbook
    wb=load_workbook(excel.files[0]['path'])
    assert wb.worksheets[0]['A2'].value=='00123'
    assert wb.worksheets[0]['A3'].value=='12345678901234567'


def test_pdf_borderless_multiple_tables(tmp_path):
    result=inspect_pdf(make_pdf(tmp_path/'wireless.pdf',borderless=True),{},tmp_path/'inspect',lambda *_:None,lambda:False)
    assert len(result.ir.tables)==2
    assert all(any(c.raw_text=='00123' for c in t.cells) for t in result.ir.tables)
    assert all(t.row_count==3 and t.column_count==3 for t in result.ir.tables)


@pytest.mark.parametrize('mode,options,expected',[('each',{},[[0],[1],[2],[3]]),('every_n',{'every_n':3},[[0,1,2],[3]]),('groups',{'groups':'3,1;4,2,2'},[[2,0],[3,1,1]]),('extract',{'groups':'3,1;4,2,2'},[[2,0,3,1,1]])])
def test_split_order_native_text(tmp_path,mode,options,expected):
    source=make_pdf(tmp_path/'pages.pdf',pages=4)
    result=convert_pdf(source,'pdf_split',{'split_mode':mode,**options},tmp_path/'out',lambda *_:None,lambda:False)
    assert len(result.files)==len(expected)
    for file,pages in zip(result.files,expected):
        reader=PdfReader(file['path'])
        assert [p.extract_text().split()[0] for p in reader.pages]==[f'PAGE-{i+1:02d}' for i in pages]


def test_visible_index_excludes_hidden_text(tmp_path):
    source=make_pdf(tmp_path/'index.pdf')
    result=convert_pdf(source,'pdf_split',{'split_mode':'crop','cuts_pt':[100]},tmp_path/'out',lambda *_:None,lambda:False)
    assert '00123' in result.ir.mappings[0]['visible_text']
    assert '00123' not in result.ir.mappings[1]['visible_text']
    # This distinction is intentional: CropBox is not redaction.
    assert '00123' in PdfReader(result.files[0]['path']).pages[1].extract_text()


def test_password_cancel_and_invalid_range(tmp_path):
    source=make_pdf(tmp_path/'plain.pdf')
    writer=PdfWriter(); writer.append(source); writer.encrypt('test-password')
    encrypted=tmp_path/'encrypted.pdf'
    with encrypted.open('wb') as f: writer.write(f)
    with pytest.raises(ToolError,match='密码'): inspect_pdf(encrypted,{},tmp_path/'bad',lambda *_:None,lambda:False)
    assert inspect_pdf(encrypted,{'password':'test-password'},tmp_path/'good',lambda *_:None,lambda:False).ir.pages
    with pytest.raises(Cancelled): convert_pdf(source,'pdf_split',{},tmp_path/'cancel',lambda *_:None,lambda:True)
    with pytest.raises(ToolError): parse_page_groups('1-8',3)
    assert parse_page_groups('2,1;2,3',3,'deduplicate')==[[1,0],[2]]


def test_long_page_suggestions_avoid_native_rows(tmp_path):
    source=make_pdf(tmp_path/'long.pdf',height=1800,table=True)
    result=split_suggestions(source,0,300)
    assert result['cuts_pt'] and all(a<b for a,b in zip(result['cuts_pt'],result['cuts_pt'][1:]))
    for cut in result['cuts_pt']:
        assert not any(p['bbox_pt'][1]+.1<cut<p['bbox_pt'][3]-.1 for p in result['protected_regions'] if p['kind']=='text')


def test_scanned_grid_reconstruction_merges_and_unknown(tmp_path):
    from PIL import Image,ImageDraw
    from app.services.document_tools.local_ocr import recover_tables
    image=Image.new('RGB',(900,500),'white'); draw=ImageDraw.Draw(image)
    for y in [50,150,250,350]: draw.line((50,y,850,y),fill='black',width=3)
    for x in [50,450,850]: draw.line((x,50,x,350),fill='black',width=3)
    draw.line((250,150,250,350),fill='black',width=3)
    draw.text((500,180),'UNREADABLE',fill='black')
    lines=[{'text':'00123','bbox_pt':[40,95,65,110],'signal_score':.97}]
    tables=recover_tables(image,lines,[0,0,450,250],0,'scan')
    assert len(tables)==1
    table=tables[0]
    assert (table.row_count,table.column_count)==(3,3)
    assert table.cells[0].colspan==2
    assert any(c.raw_text=='00123' for c in table.cells)
    assert any(c.resolution=='unknown' for c in table.cells)
    assert any(c.raw_text=='' and c.resolution=='resolved' for c in table.cells)


def test_numeric_types_preserve_ambiguous_and_precision(tmp_path):
    result=convert_pdf(make_pdf(tmp_path/'numbers.pdf',table=True),'pdf_to_excel',{'ai_mode':'off'},tmp_path/'out',lambda *_:None,lambda:False)
    from openpyxl import load_workbook
    sheet=load_workbook(result.files[0]['path']).worksheets[0]
    assert sheet['A2'].value=='00123' and sheet['A3'].value=='12345678901234567'
    assert sheet['B2'].value==12 and sheet['C2'].value==12.5
    assert sheet['C3'].value=='1,234'
    assert any(i.code=='NUMERIC_LOCALE_AMBIGUOUS' for i in result.ir.issues)


def test_merged_native_cells_preserve_grid(tmp_path):
    source=make_pdf(tmp_path/'merged.pdf')
    reader=PdfReader(source); writer=PdfWriter(); writer.append(reader)
    page=writer.pages[0]
    content='40 300 m 340 300 l S 40 330 m 340 330 l S 40 360 m 340 360 l S 40 300 m 40 360 l S 340 300 m 340 360 l S 140 300 m 140 330 l S 240 300 m 240 360 l S BT /F1 10 Tf 50 342 Td (MERGED) Tj ET BT /F1 10 Tf 50 312 Td (00123) Tj ET'
    stream=DecodedStreamObject(); stream.set_data(content.encode()); page[NameObject('/Contents')]=writer._add_object(stream)
    merged=tmp_path/'merged-grid.pdf'
    with merged.open('wb') as f: writer.write(f)
    result=inspect_pdf(merged,{},tmp_path/'inspect',lambda *_:None,lambda:False)
    assert len(result.ir.tables)==1
    table=result.ir.tables[0]
    assert table.column_count==3 and table.row_count==2
    assert table.cells[0].colspan==2 and table.cells[0].raw_text=='MERGED'


def test_low_quality_scan_is_review_candidate(tmp_path,monkeypatch):
    from PIL import Image,ImageDraw
    from app.services.document_tools import local_ocr
    # An actual raster PDF with no text layer; contract test fixes the local
    # recognizer output and does not claim measured OCR accuracy.
    image=Image.new('RGB',(300,180),'white'); ImageDraw.Draw(image).text((25,70),'00123 ?',fill='gray')
    source=tmp_path/'low-quality.pdf'; image.save(source,'PDF')
    monkeypatch.setattr(local_ocr,'recognize',lambda *_:[{'text':'00123 ?','bbox_pt':[25,70,95,85],'signal_score':.51}])
    result=convert_pdf(source,'pdf_to_word',{'ai_mode':'off'},tmp_path/'out',lambda *_:None,lambda:False)
    assert result.ir.pages[0].classification=='scanned'
    assert any(i.code=='OCR_REVIEW_REQUIRED' for i in result.ir.issues)
    assert any(b.source.method=='local_ocr' and '00123' in b.text for b in result.ir.blocks)


@pytest.mark.parametrize('variant',['clear','blur','skew','stamp','lowcontrast','small'])
def test_real_local_ocr_scan_matrix(tmp_path,variant):
    import json
    from PIL import Image,ImageDraw,ImageFont,ImageFilter
    truth='ITEM 00123 QTY 12 AMOUNT 12.50'
    image=Image.new('RGB',(1100,320),'white'); draw=ImageDraw.Draw(image)
    draw.text((35,80),truth,font=ImageFont.load_default(size=42),fill=(155,155,155) if variant=='lowcontrast' else 'black')
    if variant=='blur': image=image.filter(ImageFilter.GaussianBlur(2.2))
    if variant=='skew': image=image.rotate(3,expand=True,fillcolor='white')
    if variant=='stamp': ImageDraw.Draw(image).ellipse((140,50,320,210),outline='red',width=6)
    if variant=='small': image=image.resize((330,96))
    source=tmp_path/f'{variant}.pdf'; image.save(source,'PDF')
    result=convert_pdf(source,'pdf_to_word',{'ai_mode':'off'},tmp_path/'out',lambda *_:None,lambda:False)
    prediction=' '.join(b.text for b in result.ir.blocks if b.source.method=='local_ocr')
    # Character-level Levenshtein distance against independent known source.
    previous=list(range(len(truth)+1))
    for char in prediction:
        current=[previous[0]+1]
        for j,expected in enumerate(truth):
            current.append(min(current[-1]+1,previous[j+1]+1,previous[j]+(char!=expected)))
        previous=current
    report={'sample':variant,'engine':'RapidOCR-local-real','truth':truth,'prediction':prediction,'cer':previous[-1]/len(truth),
            'key_fields_exact':{value:value in prediction for value in ['00123','12','12.50']},'needs_review':any(i.code=='OCR_REVIEW_REQUIRED' for i in result.ir.issues),
            'automatic_pass':False,'table_structure':'not_applicable'}
    (tmp_path/'ocr-metrics.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    assert report['needs_review']
    assert result.files[0]['path'].stat().st_size>1000
    if variant=='clear': assert all(report['key_fields_exact'].values())


def test_region_reocr_preserves_manual_and_native(tmp_path,monkeypatch):
    from app.services.document_tools import local_ocr
    source=make_pdf(tmp_path/'source.pdf',table=True)
    ir=inspect_pdf(source,{},tmp_path/'inspection',lambda *_:None,lambda:False).ir
    ir.tables[0].cells[3].display_text='USER-0007'; ir.tables[0].cells[3].value='USER-0007'; ir.tables[0].cells[3].source.method='manual'
    original=ir.tables[0].cells[0].raw_text
    monkeypatch.setattr(local_ocr,'recognize',lambda *_:[{'text':'CANDIDATE-00123','bbox_pt':[50,160,120,175],'signal_score':.88}])
    result=convert_pdf(source,'pdf_to_excel',{'ai_mode':'off','_recognize_region':{'page_index':0,'bbox_pt':[40,120,340,210]}},tmp_path/'revised',lambda *_:None,lambda:False,ir=ir)
    assert result.ir.tables[0].cells[3].display_text=='USER-0007'
    assert result.ir.tables[0].cells[0].raw_text==original
    assert any('CANDIDATE-00123' in value for issue in result.ir.issues for value in issue.candidates)
    assert len(result.ir.tables)==2  # Region candidates must not duplicate native tables in export.


def test_real_scanned_merged_table_to_editable_excel(tmp_path):
    import json
    from PIL import Image,ImageDraw,ImageFont
    from openpyxl import load_workbook
    image=Image.new('RGB',(1000,440),'white'); draw=ImageDraw.Draw(image); font=ImageFont.load_default(size=34)
    for y in [40,140,240,340]: draw.line((30,y,950,y),fill='black',width=3)
    for x in [30,600,950]: draw.line((x,40,x,340),fill='black',width=3)
    draw.line((330,140,330,340),fill='black',width=3)
    values=[(50,70,'ITEM QTY'),(620,70,'AMOUNT'),(50,170,'00123'),(350,170,'12'),(620,170,'12.50'),(50,270,'00456'),(350,270,'24'),(620,270,'24.50')]
    for x,y,text in values: draw.text((x,y),text,font=font,fill='black')
    source=tmp_path/'merged-scan.pdf'; image.save(source,'PDF')
    result=convert_pdf(source,'pdf_to_excel',{'ai_mode':'off'},tmp_path/'out',lambda *_:None,lambda:False)
    assert len(result.ir.tables)==1
    table=result.ir.tables[0]
    assert (table.row_count,table.column_count)==(3,3) and table.cells[0].colspan==2
    sheet=load_workbook(result.files[0]['path']).worksheets[0]
    assert sheet['A2'].value=='00123' and sheet['B2'].value==12 and sheet['C2'].value==12.5
    assert str(next(iter(sheet.merged_cells.ranges)))=='A1:B1'
    (tmp_path/'ocr-table-metrics.json').write_text(json.dumps({'sample':'real-scanned-merged-table','table_structure_exact':True,'critical_fields_exact':True,'automatic_pass':False,
        'source_values':[v[2] for v in values],'recognized_cells':[c.raw_text for c in table.cells]},indent=2),encoding='utf-8')


def test_split_annotation_policy_external_preserved_internal_removed(tmp_path):
    from pypdf.annotations import Link,Text
    writer=PdfWriter(); writer.append(make_pdf(tmp_path/'base.pdf',pages=2))
    writer.add_annotation(0,Link(rect=(40,40,140,60),url='https://example.invalid/reference'))
    writer.add_annotation(0,Link(rect=(40,70,140,90),target_page_index=1))
    writer.add_annotation(0,Text(rect=(150,50,180,80),text='annotation evidence'))
    source=tmp_path/'annotated.pdf'
    with source.open('wb') as stream: writer.write(stream)
    result=convert_pdf(source,'pdf_split',{'split_mode':'extract','groups':'1'},tmp_path/'out',lambda *_:None,lambda:False)
    page=PdfReader(result.files[0]['path']).pages[0]
    annotations=[ref.get_object() for ref in page['/Annots']]
    assert len(annotations)==2
    assert any(a.get('/A',{}).get('/URI')=='https://example.invalid/reference' for a in annotations)
    assert any(a.get('/Contents')=='annotation evidence' for a in annotations)
    assert not any('/Dest' in a for a in annotations)


def test_pdf_continuation_merge_requires_geometry_and_context():
    from app.services.document_tools.document_ir import Block,Cell,DocumentIR,Page,SourceAnchor,Table
    from app.services.document_tools.pdf_structure import merge_continuations
    def table(page,top,bottom):
        cells=[Cell(id=f'{page}-{r}-{c}',row=r,column=c,raw_text=text,display_text=text,source=SourceAnchor(page_index=page,bbox_pt=[40+c*100,top+r*30,140+c*100,top+(r+1)*30])) for r,row in enumerate([['ITEM','QTY'],['00123','12']]) for c,text in enumerate(row)]
        return Table(id=f't{page}',row_count=2,column_count=2,header_rows=[0],source_pages=[page],cells=cells,source=SourceAnchor(page_index=page,bbox_pt=[40,top,240,bottom]))
    ir=DocumentIR(source_type='pdf',pages=[Page(page_index=i,display_page_number=i+1,width_pt=400,height_pt=600) for i in range(2)],tables=[table(0,450,570),table(1,30,100)])
    merge_continuations(ir,lambda *_:None)
    assert len(ir.tables)==1 and ir.tables[0].row_count==3
    assert ir.tables[0].cells[-1].source.page_index==1
    independent=DocumentIR(source_type='pdf',pages=ir.pages,tables=[table(0,450,570),table(1,30,100)],blocks=[Block(id='unit',text='UNIT USD',source=SourceAnchor(page_index=1,bbox_pt=[40,1,140,20]))])
    merge_continuations(independent,lambda *_:None)
    assert len(independent.tables)==2


def test_partial_digit_beside_solid_occlusion_is_unknown():
    from PIL import Image,ImageDraw
    from app.services.document_tools.local_ocr import recover_tables
    image=Image.new('RGB',(900,400),'white');draw=ImageDraw.Draw(image)
    for x in [30,300,600,870]:draw.line((x,30,x,330),fill='black',width=3)
    for y in [30,130,230,330]:draw.line((30,y,870,y),fill='black',width=3)
    draw.rectangle((650,155,800,205),fill='black')
    tables=recover_tables(image,[{'text':'1','bbox_pt':[620,165,635,195],'signal_score':.99}],[0,0,900,400],0,'occlusion')
    cell=next(c for c in tables[0].cells if c.row==1 and c.column==2)
    assert cell.raw_text=='1' and cell.display_text=='[无法辨认]' and cell.value is None
    assert cell.resolution=='unknown' and cell.candidates==['1']


def test_two_column_prose_not_invented_as_table(tmp_path):
    source=make_pdf(tmp_path/'columns.pdf',width=700)
    writer=PdfWriter();writer.append(source)
    commands=[]
    for x,prefix in [(40,'LEFT'),(390,'RIGHT')]:
        for row,word in enumerate(['First','Second','Third','Fourth']):
            commands.append(f'BT /F1 10 Tf {x} {500-row*35} Td ({prefix} {word} paragraph has ordinary prose words.) Tj ET')
    content=DecodedStreamObject();content.set_data('\n'.join(commands).encode())
    writer.pages[0][NameObject('/Contents')]=writer._add_object(content)
    with source.open('wb') as stream:writer.write(stream)
    result=inspect_pdf(source,{},tmp_path/'out',lambda *_:None,lambda:False)
    assert len(result.ir.tables)==0 and len(result.ir.blocks)==8
    assert all(b.text.startswith('LEFT') for b in result.ir.blocks[:4])
    assert all(b.text.startswith('RIGHT') for b in result.ir.blocks[4:])
