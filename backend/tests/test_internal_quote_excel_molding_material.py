import re
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.services.internal_quote_excel import _build_summary_sheet


def _build_molding_material_sheet(
    *,
    t1_rows: list[dict[str, object]],
    breakdown: dict[str, object] | None,
    expected_total: float,
):
    workbook = Workbook()
    quote = SimpleNamespace(
        product_name="啤机料价汇总测试",
        quote_no="IQ-MOLDING-MATERIAL",
        remark="",
    )
    summary: dict[str, object] = {
        "t1": t1_rows,
        "t2": [],
        "t3": [],
        "t4": [],
        "shipping_pricing": {
            "markup": "1.2",
            "active_markup_moq": "10000",
            "markup_tiers": [],
            "misc_ratio": "0.02",
            "rows": [],
        },
    }
    if breakdown is not None:
        summary["molding_material_breakdown"] = breakdown

    _build_summary_sheet(
        workbook,
        quote,
        [],
        {"fx": {"rmb_hkd": "0.85", "hkd_usd": "7.8"}},
        summary,
        {"molding_hkd": expected_total},
    )
    return workbook, workbook.active


def _find_cell(sheet, value: object):
    return next(
        cell
        for row in sheet.iter_rows()
        for cell in row
        if cell.value == value
    )


@pytest.mark.parametrize(
    ("t1_rows", "breakdown", "expected_total"),
    [
        (
            [
                {"key": "material", "value": "999"},
                {"key": "imp_mat", "value": "1"},
                {"key": "dom_mat", "value": "2"},
                # A non-molding T1 category must never leak into 料价.
                {"key": "sewing_hair", "value": "13"},
            ],
            {"total_hkd": "7.25", "imported_hkd": "2.75", "domestic_hkd": "4.5"},
            7.25,
        ),
        (
            [
                {"key": "material", "value": "8.25"},
                {"key": "imp_mat", "value": "5.25"},
                {"key": "dom_mat", "value": "3"},
            ],
            None,
            8.25,
        ),
        (
            [
                {"key": "imp_mat", "value": "1.25"},
                {"key": "dom_mat", "value": "2.75"},
            ],
            None,
            4.0,
        ),
    ],
)
def test_excel_unifies_molding_material_and_builds_tax_costs_from_detail_tax_tags(
    t1_rows: list[dict[str, object]],
    breakdown: dict[str, object] | None,
    expected_total: float,
):
    workbook, sheet = _build_molding_material_sheet(
        t1_rows=t1_rows,
        breakdown=breakdown,
        expected_total=expected_total,
    )
    try:
        summary_header = _find_cell(sheet, "旺季价")
        summary_headers = [
            sheet.cell(summary_header.row, column).value
            for column in range(3, 17)
        ]
        assert summary_headers.count("料价") == 1
        assert "进口料" not in summary_headers
        assert "国内料" not in summary_headers

        material_detail_row = next(
            row
            for row in range(1, sheet.max_row + 1)
            if sheet.cell(row, 2).value == "料价"
            and sheet.cell(row, 3).value == "料价"
        )
        assert sheet.cell(material_detail_row, 4).value == pytest.approx(expected_total)

        tax_header = _find_cell(sheet, "人民币外购件成本")
        rmb_purchase_formula = sheet.cell(tax_header.row + 1, tax_header.column).value
        tax_13_formula = sheet.cell(tax_header.row + 1, tax_header.column + 1).value
        subtotal_formula = next(
            sheet.cell(row, 4).value
            for row in range(material_detail_row, tax_header.row)
            if str(sheet.cell(row, 4).value).startswith(
                f"=SUM(D{material_detail_row}:D"
            )
        )
        detail_end_match = re.fullmatch(
            rf"=SUM\(D{material_detail_row}:D(\d+)\)", subtotal_formula
        )
        assert detail_end_match is not None
        detail_end_row = int(detail_end_match.group(1))
        detail_tax_range = f"$A${material_detail_row}:$A${detail_end_row}"
        detail_amount_range = f"$D${material_detail_row}:$D${detail_end_row}"
        assert tax_13_formula == (
            f'=SUMIF({detail_tax_range},"¥13%",{detail_amount_range})'
        )
        assert rmb_purchase_formula == (
            f'=D{tax_header.row + 1}'
            f'+SUMIF({detail_tax_range},"¥1%",{detail_amount_range})'
        )
    finally:
        workbook.close()
