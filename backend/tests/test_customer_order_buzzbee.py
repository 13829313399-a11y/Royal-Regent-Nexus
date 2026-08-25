from datetime import date
from decimal import Decimal
from io import BytesIO
import importlib
import json
from pathlib import Path
import sys
from uuid import uuid4

from fastapi.testclient import TestClient
import msoffcrypto
import openpyxl
import pytest


TEST_TMP_DIR = Path(__file__).resolve().parents[1] / ".pytest-tmp"
BACKEND_DIR = Path(__file__).resolve().parents[1]
ADMIN_TEST_PASSWORD = "AdminSeed123!"

if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))


def test_customer_order_runtime_dependencies_are_declared_for_production():
    production_requirements = {
        line.strip()
        for line in (BACKEND_DIR / "requirements.prod.txt")
        .read_text(encoding="utf-8")
        .splitlines()
        if line.strip() and not line.lstrip().startswith("#")
    }

    assert {
        "xlrd>=2.0.1,<3",
        "msoffcrypto-tool>=5.4,<6",
        "lxml>=5,<7",
        "pdfplumber>=0.11,<0.12",
    } <= production_requirements


def make_client(monkeypatch) -> TestClient:
    TEST_TMP_DIR.mkdir(exist_ok=True)
    database_url = f"sqlite:///{TEST_TMP_DIR / f'customer_order_{uuid4().hex}.db'}"
    monkeypatch.setenv("DATABASE_URL", database_url)
    monkeypatch.setenv("SEED_ADMIN_PASSWORD", ADMIN_TEST_PASSWORD)
    monkeypatch.setenv("AUTHZ_MODE", "enforce")
    monkeypatch.setenv("AUTHZ_WRITES_ENABLED", "false")

    for module_name in list(sys.modules):
        if module_name == "app" or module_name.startswith("app."):
            del sys.modules[module_name]

    main = importlib.import_module("app.main")
    return TestClient(main.app)


def ensure_user(username: str, role_id: str, department: str, factory_id: str = "huaxing") -> None:
    db_module = importlib.import_module("app.db")
    auth_models = importlib.import_module("app.models.auth")
    auth_service = importlib.import_module("app.services.auth")
    user_id = f"user-{username}"
    with db_module.SessionLocal() as db:
        if db.get(auth_models.AuthUser, user_id) is None:
            salt, password_hash = auth_service.make_password_hash("123456")
            db.add(
                auth_models.AuthUser(
                    id=user_id,
                    username=username,
                    display_name=username,
                    password_salt=salt,
                    password_hash=password_hash,
                    status="active",
                    force_password_change=0,
                    created_at=auth_service.now_text(),
                    updated_at=auth_service.now_text(),
                )
            )
        binding_id = f"{user_id}:{role_id}:{factory_id}:{department}"
        if db.get(auth_models.AuthUserRole, binding_id) is None:
            db.add(
                auth_models.AuthUserRole(
                    id=binding_id,
                    user_id=user_id,
                    role_id=role_id,
                    factory_id=factory_id,
                    department=department,
                )
            )
        db.commit()


def login(client: TestClient, username: str, role_id: str, department: str) -> dict:
    ensure_user(username, role_id, department)
    response = client.post("/api/auth/login", json={"username": username, "password": "123456"})
    assert response.status_code == 200
    return response.json()


def encrypt_xlsx(payload: bytes) -> bytes:
    workbook = msoffcrypto.OfficeFile(BytesIO(payload))
    output = BytesIO()
    workbook.encrypt("2026", output)
    return output.getvalue()


def decrypt_xlsx(payload: bytes) -> bytes:
    workbook = msoffcrypto.OfficeFile(BytesIO(payload))
    assert workbook.is_encrypted()
    workbook.load_key(password="2026")
    output = BytesIO()
    workbook.decrypt(output)
    return output.getvalue()


def build_wmc_po(
    *,
    contract_no: int = 53138,
    po_no: str = "0009382481",
    quantity: int = 3000,
) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["M3"] = date(2026, 7, 24)
    sheet["L4"] = "Contract No."
    sheet["N4"] = contract_no
    sheet["P4"] = "WMC"
    sheet["A8"] = "Date of Loading:"
    sheet["F8"] = date(2026, 9, 17)
    sheet["M10"] = "INSPECTION: 17 SEPT. 26"
    sheet["A12"] = "Our Item# :"
    sheet["F12"] = 67771
    sheet["I13"] = f"PO / BON DE COMMANDE: {po_no}"
    sheet["A14"] = "Goods:"
    sheet["C14"] = "AF DOUBLE FIRE"
    sheet["A16"] = "Quantity:"
    sheet["F16"] = quantity
    sheet["A22"] = "Shipping Carton Packing"
    sheet["F22"] = 5
    sheet["A48"] = "PACKAGING REF: 67771-05-26-WMC"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def build_tottus_po() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["Q7"] = "TOTTUS -"
    sheet["M8"] = "Contract No.:"
    sheet["O8"] = 52497
    sheet["P8"] = 7
    sheet["Q8"] = "PERU"
    sheet["A10"] = "Final Inspection Date:"
    sheet["F10"] = "TBA"
    sheet["A12"] = "Date of Loading:"
    sheet["F12"] = date(2026, 6, 12)
    sheet["A16"] = "Our Item# :"
    sheet["F16"] = 45803
    sheet["A18"] = "Goods:"
    sheet["C18"] = "BELT BLASTER"
    sheet["A20"] = "Quantity:"
    sheet["F20"] = 2800
    sheet["A26"] = "Shipping Carton Packing"
    sheet["F26"] = 4
    sheet["B40"] = "Use 45803-04-26-EN packaging"
    sheet["B44"] = "European Standard & Non-Phthalates materials is required"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def build_wmu_po() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "11019 Dept 07"
    sheet["H2"] = "Issued Date:"
    sheet["J2"] = date(2026, 8, 19)
    sheet["H4"] = "S/C No.:"
    sheet["J4"] = 53233
    sheet["N4"] = "WMU"
    sheet["A6"] = "Inspection Date:"
    sheet["C6"] = "TBA"
    sheet["A8"] = "Date of Loading:"
    sheet["C8"] = "Please see attached"
    sheet["A12"] = "Our Item# :"
    sheet["C12"] = 11019
    sheet["A15"] = "Description:"
    sheet["C15"] = "PD T-REX SQUIRTER"
    sheet["A18"] = "Quantity:"
    sheet["C18"] = 38736
    sheet["A23"] = "Outer Carton Qty:"
    sheet["C23"] = 8
    sheet["A32"] = "PO#"
    sheet["C32"] = "See attached"
    sheet["B44"] = "ASTM F963 American Standard is required"

    attached = workbook.create_sheet("PO Attached")
    attached["A1"] = "WALMART USA"
    attached["A2"] = "BB ITEM#"
    attached["B2"] = 11019
    attached["C2"] = "Walmart Item#"
    attached["D2"] = 671368029
    attached["A3"] = "Description"
    attached["B3"] = "PD T-REX SQUIRTER"
    attached["G5"] = "8pcs/ctn"
    headers = (
        "S/C NO.",
        "Walmart PO#",
        "SHIP VIA",
        "Ship Window",
        "Cancel Date",
        "Ordered Qty (PCS)",
        "Total Ctns",
        "Inspection Date",
        "Remark",
    )
    for column, header in enumerate(headers, start=1):
        attached.cell(6, column).value = header
    suborders = (
        (53233, "0105570288", "SAVANNAH", date(2026, 10, 20), date(2026, 10, 27), 3840, 480, date(2026, 9, 29)),
        (53295, "0105570350", "RIDGEVILLE", date(2026, 10, 26), date(2026, 11, 2), 4408, 551, date(2026, 10, 8)),
        (53221, "0105570276", "SUFFOLK", date(2026, 10, 26), date(2026, 11, 2), 5496, 687, date(2026, 10, 8)),
        (53273, "0105570328", "MOBILE", date(2026, 10, 30), date(2026, 11, 6), 5472, 684, date(2026, 10, 20)),
        (53218, "0105570273", "HOUSTON", date(2026, 11, 2), date(2026, 11, 9), 8184, 1023, date(2026, 10, 13)),
        (53267, "0105570322", "MIDWEST", date(2026, 11, 12), date(2026, 11, 19), 4392, 549, date(2026, 10, 22)),
        (53213, "0105570268", "EASTVALE", date(2026, 11, 14), date(2026, 11, 21), 6944, 868, date(2026, 10, 27)),
    )
    for row_number, suborder in enumerate(suborders, start=7):
        for column, value in enumerate(suborder, start=1):
            attached.cell(row_number, column).value = value
    attached["E15"] = "Total:"
    attached["F15"] = 38736
    attached["G15"] = 4842

    packaging = workbook.create_sheet("Packaging Box")
    packaging["A1"] = "Packaging Ref"
    packaging["B1"] = "11019-08-26-WMU"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def build_schedule(*, include_price: bool = True) -> bytes:
    workbook = openpyxl.Workbook()
    order = workbook.active
    order.title = "接单表"
    review = workbook.create_sheet("正单评审表")
    bullet_item = workbook.create_sheet("子弹枪ITEM表")
    water_item = workbook.create_sheet("水枪ITEM表")

    order["C3"] = "P/O#:"
    order["M3"] = "单价HK$"
    for column in range(2, 20):
        order.cell(4, column).value = ""
    order["C4"] = "0001111111"
    order["D4"] = "50000"
    order["E4"] = "WMC"
    order["F4"] = "67771"
    order["G4"] = "双管枪"
    order["H4"] = "AF DOUBLE FIRE"
    order["I4"] = 100
    order["J4"] = 5
    order["K4"] = "=I4/J4"
    order["L4"] = "加拿大标准"
    if include_price:
        order["M4"] = 23.96
    order["S4"] = date(2026, 8, 1)
    order["M5"] = "合计:HK$"
    order["O5"] = "=SUM(O4:O4)"

    review["C3"] = "PO.NO"
    review["R3"] = "单价HK$"
    for column in range(2, 24):
        review.cell(4, column).value = ""
    review["D4"] = "50000"
    review["F4"] = "67771"
    review["G4"] = "双管枪"
    review["R5"] = "子弹枪合计"
    review["T5"] = "=SUM(T4:T4)"
    review["R7"] = "水枪合计"
    review["T7"] = "=SUM(T6:T6)"

    item_headers = (
        "客出单日期",
        "预备单号（OQF NO）",
        "PO.NO",
        "客名/国家",
        "产品编号",
        "中文名称",
        "产品名称",
        "数量",
        "装箱",
        "彩盒条形码",
        "包装",
        "系统",
        "行验",
        "客验",
        "客要求走货期",
        "备注",
        "客人PO",
    )
    for sheet in (bullet_item, water_item):
        for column, header in enumerate(item_headers, 1):
            sheet.cell(3, column).value = header

    bullet_item["B4"] = "2026-6-23-67771"
    bullet_item["E4"] = "67771-2"
    bullet_item["F4"] = "双管枪"
    bullet_item["G4"] = "AF DOUBLE FIRE"
    bullet_item["H4"] = 5000
    bullet_item["I4"] = 5
    bullet_item["K4"] = "67771-05-24-WMC(r1)"
    bullet_item["L4"] = "系统"
    bullet_item["M4"] = "备料单"
    bullet_item["A5"] = date(2026, 7, 1)
    bullet_item["B5"] = "2026-7-1-40210"
    bullet_item["C5"] = "53099"
    bullet_item["D5"] = "WMC"
    bullet_item["E5"] = "40210-1"
    bullet_item["F5"] = "转盘枪"
    bullet_item["G5"] = "MAYHEM OUTRAGE"
    bullet_item["H5"] = 48
    bullet_item["I5"] = 4
    bullet_item["K5"] = "40210-03-26-WM"
    bullet_item["O5"] = date(2026, 8, 10)
    bullet_item["Q5"] = "0001111111"
    bullet_item["A7"] = "已走货子弹枪"

    water_item["B4"] = "2026-1-2-11580-1/-2"
    water_item["E4"] = "11580-1/-2"
    water_item["F4"] = "鱼缸水枪"
    water_item["G4"] = "OCEAN OUTLAW BLASTER"
    water_item["H4"] = 2352
    water_item["K4"] = "11580-10-25-EN"
    water_item["L4"] = "备料单"
    water_item["A5"] = date(2026, 7, 1)
    water_item["B5"] = "2026-7-1-19570"
    water_item["C5"] = "53098"
    water_item["D5"] = "AAFES"
    water_item["E5"] = "19570-3"
    water_item["F5"] = "中水枪"
    water_item["G5"] = "GARGOYLE 2 PACK"
    water_item["H5"] = 120
    water_item["I5"] = 6
    water_item["K5"] = "19570-10-23-EN"
    water_item["O5"] = date(2026, 8, 10)
    water_item["A7"] = "已走货水枪"

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return encrypt_xlsx(output.getvalue())


def build_wmu_schedule() -> bytes:
    workbook = openpyxl.load_workbook(BytesIO(decrypt_xlsx(build_schedule())))
    order = workbook["接单表"]
    order["F4"] = "11019"
    order["G4"] = "恐龙水枪"
    order["H4"] = "PD T-REX SQUIRTER"
    order["J4"] = 8
    order["L4"] = "美国标准"
    order["M4"] = 18.5
    review = workbook["正单评审表"]
    review["F4"] = "11019"
    review["G4"] = "恐龙水枪"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return encrypt_xlsx(output.getvalue())


def upload_preview(client: TestClient):
    return client.post(
        "/api/customer-orders/buzzbee/preview",
        data={"factory_id": "huaxing", "received_date": "2026-07-27"},
        files={
            "po_file": (
                "WM-67771-53138-WMC.xlsx",
                build_wmc_po(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
            "schedule_file": (
                "2026年 BUZZ BEE 生产排期表.xls.xlsx",
                build_schedule(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
    )


def test_buzzbee_preview_requires_sales_permission_and_maps_seventeen_fields(monkeypatch):
    with make_client(monkeypatch) as client:
        assert upload_preview(client).status_code == 401

        login(client, "customer_order_engineer", "engineer", "engineering")
        assert upload_preview(client).status_code == 403

        client.post("/api/auth/logout")
        profile = login(client, "customer_order_sales", "sales_customer_owner", "sales-business")
        assert "customer_order:read" in profile["permissions"]
        assert "customer_order:export" in profile["permissions"]
        assert "customer_order:duplicate_confirm" in profile["permissions"]
        assert "customer_order:audit_read" not in profile["permissions"]

        response = upload_preview(client)
        assert response.status_code == 200, response.text
        preview = response.json()
        assert preview["summary"] == {"total": 1, "valid": 1, "warning": 0, "blocked": 0}
        assert preview["input_template"] == "BUZZBEE_WALMART_WMC_INLINE_V1"
        assert preview["target_template"] == "BUZZBEE_PRODUCTION_SCHEDULE_V1"
        assert preview["po_file_count"] == 1
        assert preview["po_file_names"] == ["WM-67771-53138-WMC.xlsx"]
        assert preview["output_file_name"] == "2026年 BUZZ BEE 生产排期表.xls.xlsx"
        row = preview["rows"][0]
        assert {
            "received_date": row["received_date"],
            "po_no": row["po_no"],
            "contract_no": row["contract_no"],
            "customer_country": row["customer_country"],
            "product_no": row["product_no"],
            "product_name_zh": row["product_name_zh"],
            "product_name_en": row["product_name_en"],
            "quantity": row["quantity"],
            "units_per_carton": row["units_per_carton"],
            "carton_count": row["carton_count"],
            "standard": row["standard"],
            "unit_price_hkd": row["unit_price_hkd"],
            "amount_hkd": row["amount_hkd"],
            "packaging": row["packaging"],
            "line_q": row["line_q"],
            "customer_q": row["customer_q"],
            "requested_ship_date": row["requested_ship_date"],
        } == {
            "received_date": "2026-07-27",
            "po_no": "0009382481",
            "contract_no": "53138",
            "customer_country": "WMC / 加拿大",
            "product_no": "67771",
            "product_name_zh": "双管枪",
            "product_name_en": "AF DOUBLE FIRE",
            "quantity": "3000",
            "units_per_carton": "5",
            "carton_count": "600",
            "standard": "加拿大标准",
            "unit_price_hkd": "23.96",
            "amount_hkd": "71880",
            "packaging": "67771-05-26-WMC",
            "line_q": "2026-09-17",
            "customer_q": "2026-09-17",
            "requested_ship_date": "2026-09-17",
        }
        assert len(row["lineage"]) == 17
        assert row["item_sheet_name"] == "子弹枪ITEM表"


def test_wmu_mainland_preview_and_export_use_po_attached_suborders(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "customer_order_wmu_sales", "sales_customer_owner", "sales-business")
        common = {"factory_id": "huaxing", "received_date": "2026-08-19"}
        po_content = build_wmu_po()
        schedule_content = build_wmu_schedule()

        def files():
            return {
                "po_file": (
                    "WMU 11019 (DEPT 07) 40 - 53284 WH.xlsx",
                    po_content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
                "schedule_file": (
                    "2026年 BUZZ BEE 生产排期表.xls.xlsx",
                    schedule_content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            }

        response = client.post(
            "/api/customer-orders/buzzbee/preview",
            data=common,
            files=files(),
        )

        assert response.status_code == 200, response.text
        preview = response.json()
        assert preview["summary"] == {"total": 7, "valid": 7, "warning": 0, "blocked": 0}
        assert preview["input_template"] == "BUZZBEE_WALMART_WMU_ATTACHED_V1"
        assert [row["contract_no"] for row in preview["rows"]] == [
            "53233",
            "53295",
            "53221",
            "53273",
            "53218",
            "53267",
            "53213",
        ]
        assert [row["po_no"] for row in preview["rows"]] == [
            "0105570288",
            "0105570350",
            "0105570276",
            "0105570328",
            "0105570273",
            "0105570322",
            "0105570268",
        ]
        first = preview["rows"][0]
        assert first["customer_country"] == "WMU / 美国"
        assert first["quantity"] == "3840"
        assert first["units_per_carton"] == "8"
        assert first["carton_count"] == "480"
        assert first["requested_ship_date"] == "2026-10-20"
        assert first["lineage"]["contract_no"] == "PO Attached!A7"
        assert first["lineage"]["po_no"] == "PO Attached!B7"
        assert first["source_po_file_name"] == "WMU 11019 (DEPT 07) 40 - 53284 WH.xlsx"

        exported = client.post(
            "/api/customer-orders/buzzbee/export",
            data={
                **common,
                "confirmed": "true",
                "preview_fingerprint": preview["preview_fingerprint"],
            },
            files=files(),
        )
        assert exported.status_code == 200, exported.text
        service = importlib.import_module("app.services.customer_order_buzzbee")
        output = service.OoxmlSchedule(decrypt_xlsx(exported.content))
        order_rows = output.read_rows("接单表")
        exported_po_numbers = {
            values.get("C")
            for values in order_rows.values()
            if values.get("D") in {"53233", "53295", "53221", "53273", "53218", "53267", "53213"}
        }
        assert exported_po_numbers == {
            "0105570288",
            "0105570350",
            "0105570276",
            "0105570328",
            "0105570273",
            "0105570322",
            "0105570268",
        }


def test_buzzbee_export_requires_confirmation_and_returns_updated_encrypted_schedule(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "customer_order_exporter", "sales_customer_owner", "sales-business")
        common = {
            "factory_id": "huaxing",
            "received_date": "2026-07-27",
        }
        files = {
            "po_file": (
                "WM-67771-53138-WMC.xlsx",
                build_wmc_po(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
            "schedule_file": (
                "2026年 BUZZ BEE 生产排期表.xls.xlsx",
                build_schedule(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        }
        rejected = client.post(
            "/api/customer-orders/buzzbee/export",
            data={**common, "confirmed": "false"},
            files=files,
        )
        assert rejected.status_code == 400

        preview_response = client.post(
            "/api/customer-orders/buzzbee/preview",
            data=common,
            files=files,
        )
        assert preview_response.status_code == 200, preview_response.text
        preview_fingerprint = preview_response.json()["preview_fingerprint"]

        exported = client.post(
            "/api/customer-orders/buzzbee/export",
            data={
                **common,
                "confirmed": "true",
                "preview_fingerprint": preview_fingerprint,
            },
            files=files,
        )
        assert exported.status_code == 200, exported.text
        assert exported.content.startswith(b"\xd0\xcf\x11\xe0")
        assert exported.headers["x-workbook-password-required"] == "true"
        assert exported.headers["x-output-template"] == "BUZZBEE_PRODUCTION_SCHEDULE_V1"
        assert exported.headers["x-preview-fingerprint"] == preview_fingerprint
        assert exported.headers["x-export-audit-id"].startswith("customer-order-export-")
        assert len(exported.headers["x-content-sha256"]) == 64
        assert "2026%E5%B9%B4%20BUZZ%20BEE" in exported.headers["content-disposition"]

        assert client.get("/api/customer-orders/audits?factory_id=huaxing").status_code == 403
        client.post("/api/auth/logout")
        supervisor_profile = login(
            client,
            "customer_order_auditor",
            "sales_customer_supervisor",
            "sales-business",
        )
        assert "customer_order:audit_read" in supervisor_profile["permissions"]
        audits_response = client.get("/api/customer-orders/audits?factory_id=huaxing")
        assert audits_response.status_code == 200, audits_response.text
        audits = audits_response.json()
        assert len(audits) == 1
        assert audits[0]["id"] == exported.headers["x-export-audit-id"]
        assert audits[0]["actor_username"] == "customer_order_exporter"
        assert audits[0]["preview_fingerprint"] == preview_fingerprint
        assert audits[0]["confirmed_issue_count"] == 0
        assert audits[0]["output_sha256"] == exported.headers["x-content-sha256"]

        service = importlib.import_module("app.services.customer_order_buzzbee")
        workbook = service.OoxmlSchedule(decrypt_xlsx(exported.content))
        order_rows = workbook.read_rows("接单表")
        review_rows = workbook.read_rows("正单评审表")
        item_rows = workbook.read_rows("子弹枪ITEM表")
        assert order_rows[5]["C"] == "0009382481"
        assert order_rows[5]["D"] == "53138"
        assert order_rows[5]["F"] == "67771"
        assert order_rows[5]["K"] == "600"
        assert order_rows[5]["O"] == "71880"
        assert review_rows[5]["C"] == "0009382481"
        assert review_rows[5]["D"] == "53138"
        assert review_rows[5]["F"] == "67771"
        assert review_rows[5]["K"] == "600"
        assert review_rows[5]["T"] == "71880"
        assert order_rows[6]["M"] == "合计:HK$"
        assert review_rows[6]["R"] == "子弹枪合计"
        assert workbook.read_cell_formula("接单表", 6, "O") == "SUM(O4:O5)"
        assert workbook.read_cell_formula("正单评审表", 6, "T") == "SUM(T4:T5)"
        assert item_rows[4]["B"] == "2026-6-23-67771"
        assert item_rows[4]["C"] == "53138"
        assert item_rows[4]["D"] == "WMC"
        assert item_rows[4]["E"] == "67771-2"
        assert item_rows[4]["H"] == "3000"
        assert item_rows[4]["I"] == "5"
        assert item_rows[4]["K"] == "67771-05-26-WMC"
        assert item_rows[4]["L"] == "系统"
        assert item_rows[4]["M"] == item_rows[4]["N"] == item_rows[4]["O"]
        assert item_rows[4]["Q"] == "0009382481"
        assert item_rows[5]["H"] == "2000"
        assert item_rows[5]["M"] == "备料单"
        assert workbook.read_cell_formula("子弹枪ITEM表", 5, "H") == "5000-3000"


def test_batch_preview_and_export_merge_multiple_po_files_without_renaming_schedule(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "customer_order_batch", "sales_customer_owner", "sales-business")
        common = {
            "factory_id": "huaxing",
            "received_date": "2026-07-27",
        }
        first_po_content = build_wmc_po()
        second_po_content = build_wmc_po(contract_no=53139, po_no="0009382482")
        schedule_content = build_schedule()

        def batch_files():
            return [
                (
                    "po_files",
                    (
                        "WM-67771-53138-WMC.xlsx",
                        first_po_content,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                ),
                (
                    "po_files",
                    (
                        "WM-67771-53139-WMC.xlsx",
                        second_po_content,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                ),
                (
                    "schedule_file",
                    (
                        "2026年 BUZZ BEE 生产排期表.xls.xlsx",
                        schedule_content,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                ),
            ]

        preview_response = client.post(
            "/api/customer-orders/buzzbee/preview-batch",
            data=common,
            files=batch_files(),
        )
        assert preview_response.status_code == 200, preview_response.text
        preview = preview_response.json()
        assert preview["po_file_count"] == 2
        assert preview["po_file_names"] == [
            "WM-67771-53138-WMC.xlsx",
            "WM-67771-53139-WMC.xlsx",
        ]
        assert preview["summary"] == {"total": 2, "valid": 2, "warning": 0, "blocked": 0}
        assert preview["output_file_name"] == "2026年 BUZZ BEE 生产排期表.xls.xlsx"
        assert [row["source_po_file_name"] for row in preview["rows"]] == preview["po_file_names"]

        exported = client.post(
            "/api/customer-orders/buzzbee/export-batch",
            data={
                **common,
                "confirmed": "true",
                "skipped_issue_keys": "[]",
                "preview_fingerprint": preview["preview_fingerprint"],
            },
            files=batch_files(),
        )
        assert exported.status_code == 200, exported.text
        assert exported.headers["x-po-file-count"] == "2"
        disposition = exported.headers["content-disposition"]
        assert "%E5%B7%B2%E5%A1%AB" not in disposition
        assert disposition.endswith(
            "2026%E5%B9%B4%20BUZZ%20BEE%20%E7%94%9F%E4%BA%A7%E6%8E%92%E6%9C%9F%E8%A1%A8.xls.xlsx"
        )

        service = importlib.import_module("app.services.customer_order_buzzbee")
        workbook = service.OoxmlSchedule(decrypt_xlsx(exported.content))
        order_rows = workbook.read_rows("接单表")
        review_rows = workbook.read_rows("正单评审表")
        item_rows = workbook.read_rows("子弹枪ITEM表")
        assert [order_rows[row]["D"] for row in (5, 6)] == ["53138", "53139"]
        assert [review_rows[row]["D"] for row in (5, 6)] == ["53138", "53139"]
        assert [item_rows[row]["C"] for row in (4, 5)] == ["53138", "53139"]
        assert [item_rows[row]["Q"] for row in (4, 5)] == ["0009382481", "0009382482"]
        assert item_rows[6]["H"] == "-1000"
        assert item_rows[6]["M"] == "备料单"
        assert workbook.read_cell_formula("子弹枪ITEM表", 6, "H") == "5000-3000-3000"
        assert order_rows[7]["M"] == "合计:HK$"
        assert review_rows[7]["R"] == "子弹枪合计"
        assert workbook.read_cell_formula("接单表", 7, "O") == "SUM(O4:O6)"
        assert workbook.read_cell_formula("正单评审表", 7, "T") == "SUM(T4:T6)"


def test_missing_price_can_be_explicitly_skipped_and_exports_blank_amounts(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "customer_order_price_skip", "sales_customer_owner", "sales-business")
        common = {
            "factory_id": "huaxing",
            "received_date": "2026-07-27",
        }
        po_content = build_wmc_po()
        schedule_content = build_schedule(include_price=False)

        def files():
            return {
                "po_file": (
                    "WM-67771-53138-WMC.xlsx",
                    po_content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
                "schedule_file": (
                    "2026年 BUZZ BEE 生产排期表.xls.xlsx",
                    schedule_content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            }

        preview_response = client.post(
            "/api/customer-orders/buzzbee/preview",
            data=common,
            files=files(),
        )
        assert preview_response.status_code == 200, preview_response.text
        preview = preview_response.json()
        assert preview["summary"] == {"total": 1, "valid": 0, "warning": 0, "blocked": 1}
        price_issue = next(
            issue
            for issue in preview["rows"][0]["issues"]
            if issue["code"] == "missing_unit_price"
        )
        assert price_issue["can_skip"] is True
        assert price_issue["skip_label"] == "单价及金额留空，稍后由跟客补充"
        assert price_issue["skip_key"]

        rejected = client.post(
            "/api/customer-orders/buzzbee/export",
            data={
                **common,
                "confirmed": "true",
                "skipped_issue_keys": "[]",
                "preview_fingerprint": preview["preview_fingerprint"],
            },
            files=files(),
        )
        assert rejected.status_code == 400
        assert "没有可确认的 HKD 单价" in rejected.json()["detail"]

        exported = client.post(
            "/api/customer-orders/buzzbee/export",
            data={
                **common,
                "confirmed": "true",
                "skipped_issue_keys": json.dumps([price_issue["skip_key"]]),
                "preview_fingerprint": preview["preview_fingerprint"],
                "confirmation_reason": "测试阶段核对缺失单价处理",
            },
            files=files(),
        )
        assert exported.status_code == 200, exported.text
        assert exported.headers["x-skipped-issue-count"] == "1"

        service = importlib.import_module("app.services.customer_order_buzzbee")
        workbook = service.OoxmlSchedule(decrypt_xlsx(exported.content))
        order_rows = workbook.read_rows("接单表")
        review_rows = workbook.read_rows("正单评审表")
        assert order_rows[5].get("M") in (None, "")
        assert order_rows[5].get("O") in (None, "")
        assert review_rows[5].get("R") in (None, "")
        assert review_rows[5].get("T") in (None, "")


def test_item_export_routes_water_product_and_deducts_preparation_stock():
    service = importlib.import_module("app.services.customer_order_buzzbee")
    plain_schedule, _ = service._decrypt_schedule(build_schedule())
    workbook = service.OoxmlSchedule(plain_schedule)
    result = service._write_item_order_row(
        workbook,
        {
            "received_date": "2026-07-27",
            "po_no": "",
            "contract_no": "52335",
            "customer_name": "AAFES",
            "product_no": "11580",
            "product_name_zh": "鱼缸水枪",
            "product_name_en": "OCEAN OUTLAW BLASTER",
            "quantity": "180",
            "units_per_carton": "6",
            "packaging": "11580-10-25-EN",
            "line_q": "2026-01-28",
            "customer_q": "To be Advised",
            "requested_ship_date": "2026-02-02",
        },
    )
    item_rows = workbook.read_rows("水枪ITEM表")
    assert result == {
        "sheet_name": "水枪ITEM表",
        "inserted_row": 4,
        "item_no": "11580-1/-2",
        "oqf_no": "2026-1-2-11580-1/-2",
    }
    assert item_rows[4]["C"] == "52335"
    assert item_rows[4]["D"] == "AAFES"
    assert item_rows[4]["E"] == "11580-1/-2"
    assert item_rows[4]["H"] == "180"
    assert item_rows[4]["I"] == "6"
    assert item_rows[4]["L"] == "系统"
    assert item_rows[4]["M"] == str((date(2026, 1, 28) - date(1899, 12, 30)).days)
    assert item_rows[4]["N"] == "To be Advised"
    assert item_rows[4]["O"] == str((date(2026, 2, 2) - date(1899, 12, 30)).days)
    assert item_rows[4].get("Q") in (None, "")
    assert item_rows[5]["H"] == "2172"
    assert item_rows[5]["L"] == "备料单"
    assert workbook.read_cell_formula("水枪ITEM表", 5, "H") == "2352-180"


def test_ordinary_xlsx_contract_uses_labels_instead_of_aafes_or_wmc_gate():
    service = importlib.import_module("app.services.customer_order_buzzbee")

    parsed = service.parse_po(
        "TOTTUS SM 52497 - 45803 WH 7.xlsx",
        build_tottus_po(),
    )[0]

    assert parsed.input_template == service.STANDARD_TEMPLATE
    assert parsed.values == {
        "contract_no": "52497",
        "inspection_date": "",
        "requested_ship_date": "2026-06-12",
        "product_no": "45803",
        "product_name_en": "BELT BLASTER",
        "quantity": 2800,
        "units_per_carton": 4,
        "po_no": "",
        "customer_name": "TOTTUS -PERU",
        "country": "秘鲁",
        "standard": "欧洲标准",
        "packaging": "45803-04-26-EN",
        "inspection_raw": "TBA",
    }
    assert parsed.lineage["contract_no"] == "Sheet1!O8"
    assert parsed.lineage["customer_name"] == "Sheet1!Q7 / Sheet1!Q8"
    assert parsed.lineage["standard"] == "合同条款 · European Standard"


def test_aafe_filename_typo_still_uses_aafes_customer_profile():
    service = importlib.import_module("app.services.customer_order_buzzbee")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["M8"] = "Contract No.:"
    sheet["O8"] = 53080
    sheet["A12"] = "Date of Loading:"
    sheet["F12"] = date(2026, 10, 1)
    sheet["A16"] = "Our Item# :"
    sheet["F16"] = 40210
    sheet["A18"] = "Goods:"
    sheet["C18"] = "MAYHEM OUTRAGE"
    sheet["A20"] = "Quantity:"
    sheet["F20"] = 80
    sheet["A26"] = "Shipping Carton Packing"
    sheet["F26"] = 4
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    parsed = service.parse_po("AAAFE SC# 53080 - 40210 WH 2 Rev 1.xlsx", output.getvalue())[0]

    assert parsed.values["customer_name"] == "AAFES"
    assert parsed.values["country"] == "美国"
    assert parsed.values["standard"] == "美国标准"


def test_wmu_mainland_xlsx_contract_expands_po_attached_suborders():
    service = importlib.import_module("app.services.customer_order_buzzbee")

    parsed = service.parse_po(
        "WMU 11019 (DEPT 07) 40 - 53284 WH.xlsx",
        build_wmu_po(),
    )

    assert len(parsed) == 7
    assert {line.input_template for line in parsed} == {service.WMU_TEMPLATE}
    assert all(line.values["product_no"] == "11019" for line in parsed)
    assert all(line.values["product_name_en"] == "PD T-REX SQUIRTER" for line in parsed)
    assert all(line.values["units_per_carton"] == 8 for line in parsed)
    assert all(line.values["customer_name"] == "WMU" for line in parsed)
    assert sum(line.values["quantity"] for line in parsed) == 38736
    assert sum(line.values["declared_carton_count"] for line in parsed) == 4842

    first = parsed[0]
    assert first.values == {
        "contract_no": "53233",
        "inspection_date": "2026-09-29",
        "requested_ship_date": "2026-10-20",
        "product_no": "11019",
        "product_name_en": "PD T-REX SQUIRTER",
        "quantity": 3840,
        "units_per_carton": 8,
        "declared_carton_count": 480,
        "po_no": "0105570288",
        "customer_name": "WMU",
        "country": "美国",
        "standard": "美国标准",
        "packaging": "11019-08-26-WMU",
        "inspection_raw": "2026-09-29 00:00:00",
    }
    assert first.lineage["contract_no"] == "PO Attached!A7"
    assert first.lineage["po_no"] == "PO Attached!B7"
    assert first.lineage["quantity"] == "PO Attached!F7"
    assert first.lineage["requested_ship_date"] == "PO Attached!D7"
    assert first.lineage["inspection_date"] == "PO Attached!H7"
    assert first.lineage["standard"] == "合同条款 · U.S. Standard"

    last = parsed[-1]
    assert last.values["contract_no"] == "53213"
    assert last.values["po_no"] == "0105570268"
    assert last.values["quantity"] == 6944
    assert last.values["requested_ship_date"] == "2026-11-14"
    assert last.values["inspection_date"] == "2026-10-27"


def test_wmu_mainland_contract_requires_po_attached_sheet():
    service = importlib.import_module("app.services.customer_order_buzzbee")
    workbook = openpyxl.load_workbook(BytesIO(build_wmu_po()))
    del workbook["PO Attached"]
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    with pytest.raises(service.CustomerOrderWorkbookError, match="缺少 PO Attached 子订单页"):
        service.parse_po(
            "WMU 11019 (DEPT 07) 40 - 53284 WH.xlsx",
            output.getvalue(),
        )


def test_wmu_declared_carton_conflict_is_blocked():
    service = importlib.import_module("app.services.customer_order_buzzbee")
    workbook = openpyxl.load_workbook(BytesIO(build_wmu_po()))
    workbook["PO Attached"]["G7"] = 481
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    parsed = service.parse_po(
        "WMU 11019 (DEPT 07) 40 - 53284 WH.xlsx",
        output.getvalue(),
    )
    preview = service._build_preview_rows(
        parsed,
        {
            "11019": service.ScheduleLookup(
                product_name_zh="恐龙水枪",
                unit_price_hkd=Decimal("18.5"),
                existing_rows=[],
            )
        },
        "2026-08-19",
    )

    assert preview[0]["status"] == "blocked"
    issue = next(item for item in preview[0]["issues"] if item["code"] == "carton_count_conflict")
    assert issue["message"] == "PO Attached 标示 481 箱，但数量 ÷ 装箱数为 480 箱"


@pytest.mark.parametrize("marker", ["Indonesia", "印尼"])
def test_explicit_indonesia_xlsx_contract_remains_out_of_scope(marker):
    service = importlib.import_module("app.services.customer_order_buzzbee")

    with pytest.raises(service.CustomerOrderWorkbookError, match="印尼合同"):
        service.parse_po(f"{marker} schedule.xlsx", build_tottus_po())


def test_po_number_is_optional_for_standard_customer_but_required_for_walmart():
    service = importlib.import_module("app.services.customer_order_buzzbee")
    base_values = {
        "po_no": "",
        "contract_no": "52496",
        "customer_name": "AAFES",
        "country": "美国",
        "product_no": "40210",
        "product_name_en": "MAYHEM OUTRAGE",
        "quantity": "48",
        "units_per_carton": "4",
        "requested_ship_date": "2026-08-10",
        "inspection_date": "2026-08-10",
        "standard": "美国标准",
        "packaging": "40210-01-26-EN",
    }
    schedule_index = {
        "40210": service.ScheduleLookup(
            product_name_zh="转盘枪",
            unit_price_hkd=Decimal("9.25"),
            existing_rows=[],
        )
    }
    standard = service.ParsedPoLine(
        values=base_values,
        lineage={},
        input_template=service.STANDARD_TEMPLATE,
    )
    ordinary_row = service._build_preview_rows(
        [standard],
        schedule_index,
        "2026-07-27",
    )[0]
    assert ordinary_row["po_no"] == ""
    assert ordinary_row["status"] == "valid"
    assert not any(issue["field"] == "po_no" for issue in ordinary_row["issues"])

    walmart = service.ParsedPoLine(
        values={**base_values, "customer_name": "WMC", "country": "加拿大"},
        lineage={},
        input_template=service.WMC_TEMPLATE,
    )
    walmart_row = service._build_preview_rows(
        [walmart],
        schedule_index,
        "2026-07-27",
    )[0]
    po_issue = next(issue for issue in walmart_row["issues"] if issue["field"] == "po_no")
    assert walmart_row["status"] == "blocked"
    assert po_issue["message"] == "WM 客 P/O# 未能从 PO 提取"
    assert po_issue["can_skip"] is False


def test_batch_duplicate_order_line_is_test_stage_confirmable():
    service = importlib.import_module("app.services.customer_order_buzzbee")
    base = {
        "po_no": "PO-1",
        "contract_no": "SC-1",
        "product_no": "ITEM-1",
        "requested_ship_date": "2026-08-20",
        "status": "valid",
        "status_label": "有效",
        "issues": [],
    }
    rows = [
        {**base, "id": "row-1", "source_po_file_name": "first.xlsx"},
        {**base, "id": "row-2", "source_po_file_name": "second.xlsx", "issues": []},
    ]

    service._mark_batch_duplicates(rows)

    issue = rows[1]["issues"][0]
    assert rows[1]["status"] == "blocked"
    assert issue["code"] == "duplicate_batch_order_line"
    assert issue["can_skip"] is True
    assert issue["skip_key"] == "row-2|duplicate_batch_order_line|po_no"
