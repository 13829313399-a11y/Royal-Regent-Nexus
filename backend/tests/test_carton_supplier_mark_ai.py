import hashlib
import io
import json
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from openpyxl import Workbook
from pydantic import ValidationError

from app.services import carton_mark_ai_pilot as pilot
from app.services import carton_supplier_mark_ai as ai
from app.schemas.carton_supplier_portal import SupplierMarkPdfRequest


def original(filled=False, **changes):
    book=Workbook(); sheet=book.active;sheet.title="SM"
    for row,text in enumerate(("ZUMIEZ INC.","PO NUMBER:00484860","ZUMIEZ ITEM NO.: 00415647",
            "ITEM DESCRIPTION: CURRENT PRODUCT","QUANTITY:","CTN QTY:","COUNTRY OF ORIGIN: CHINA"),1):
        sheet.cell(row,1,text)
    if filled:
        sheet.cell(5,2,6);sheet.cell(5,3,"PCS")
        sheet.cell(6,2,"/");sheet.cell(6,3,9);sheet.cell(6,4,"CTNS")
    stream=io.BytesIO();book.save(stream);book.close()
    template=pilot.load_template()
    result=dict(excel_file_name="current.xlsx",excel_bytes=stream.getvalue(),contract_no="4500999001",item_no="INTERNAL-001",
        issue_id="issue-current",supplements={},order_fields={"quantity":"7","carton_total":"11"},source_item="",
        provider_signature="signature",ai_consent_version=ai.CONSENT_VERSION,
        layout=dict(config=template,reference_bytes=b"fixed artwork only"))
    result.update(changes);return result


def plan_from(entries,native,identity,fallbacks):
    fields={}
    for field in pilot.FIELDS:
        if field in identity: source_id,value="ORDER",identity[field]
        elif field in native:
            value=native[field]
            source_id=next(e.id for e in entries if e.kind=="current_excel" and value in e.text)
        else:source_id,value="FILL-"+field,fallbacks[field]
        fields[field]=dict(source_id=source_id,value=value)
    return json.dumps(dict(template_id="zumiez-review-v1",layout_variant="reference",fields=fields))


def test_current_order_identity_and_counts_never_come_from_historical_pdf():
    entries,native,identity,fallbacks=ai.current_evidence(original())
    assert identity==dict(contract="4500999001",item="INTERNAL-001")
    assert fallbacks==dict(quantity="7 PCS",carton_total="11")
    assert not any(e.kind=="reference_pdf" for e in entries)
    raw=plan_from(entries,native,identity,fallbacks)
    _,values,proof=pilot.validate_plan(raw,entries,pilot.load_template(),native,identity=identity,fallbacks=fallbacks)
    assert values["customer_po"]=="00484860" and values["article"]=="00415647"
    assert proof["contract"]["kind"]=="issued_order" and values["carton_total"]=="11"


def test_populated_excel_counts_take_precedence_and_separate_cells_are_cited_as_one_row():
    entries,native,identity,fallbacks=ai.current_evidence(original(filled=True))
    assert native["quantity"]=="6 PCS" and native["carton_total"]=="9" and not fallbacks
    raw=plan_from(entries,native,identity,fallbacks)
    _,values,proof=pilot.validate_plan(raw,entries,pilot.load_template(),native,identity=identity,fallbacks=fallbacks)
    assert values["quantity"]=="6 PCS" and proof["quantity"]["kind"]=="current_excel"
    with pytest.raises(pilot.PilotError,match="不能覆盖"):
        ai.current_evidence(original(filled=True,supplements={"quantity":"99 PCS"}))


def test_missing_counts_block_before_provider_and_manual_supplements_are_order_only():
    with pytest.raises(pilot.PilotError,match="缺少"):
        ai.current_evidence(original(order_fields={}))
    entries,_,_,fallbacks=ai.current_evidence(original(supplements={"quantity":"8 PCS","carton_total":"13"}))
    assert fallbacks==dict(quantity="8 PCS",carton_total="13")
    assert all(e.kind=="order_supplement" for e in entries if e.id.startswith("FILL"))


def test_runtime_rejects_a_historical_value_even_when_it_has_a_valid_text_reference():
    entries,native,identity,fallbacks=ai.current_evidence(original())
    raw=json.loads(plan_from(entries,native,identity,fallbacks))
    entries.append(pilot.Evidence(id="OLD",kind="reference_pdf",location="old:1",text="42"))
    raw["fields"]["carton_total"]={"source_id":"OLD","value":"42"}
    with pytest.raises(pilot.PilotError):
        pilot.validate_plan(json.dumps(raw),entries,pilot.load_template(),native,identity=identity,fallbacks=fallbacks)


def test_prepare_sends_only_current_fields_and_preserves_provider_and_source_provenance(monkeypatch):
    src=original();entries,native,identity,fallbacks=ai.current_evidence(src);calls=[]
    monkeypatch.setattr(ai,"configuration",lambda:({},pilot.load_template(),b"fixed artwork only","signature"))
    def request(sent,template,connection,**kwargs):
        calls.append(sent)
        assert "issued_order" in kwargs["prompt"]
        assert all(e.kind!="reference_pdf" for e in sent)
        return plan_from(entries,native,identity,fallbacks),{"model":"fake-model","usage":{"total_tokens":50}}
    monkeypatch.setattr(pilot,"request_plan",request)
    monkeypatch.setattr(ai,"_render",lambda original,template,values,proof,variant:(b"new pdf",{"pages":1}))
    content,pages,notes,metadata=ai.prepare(src)
    assert content==b"new pdf" and pages==1 and len(calls)==1
    assert metadata["field_sources"]["item"]["kind"]=="issued_order"
    assert metadata["consent_version"]==ai.CONSENT_VERSION and "value" not in metadata["field_sources"]["item"]
    assert not any("42" in e.text or "12 PCS" in e.text for e in calls[0])


def test_changed_configuration_and_missing_counts_do_not_call_provider_and_capacity_is_released(monkeypatch):
    calls=[]
    monkeypatch.setattr(pilot,"request_plan",lambda *a,**kw:calls.append(a))
    monkeypatch.setattr(ai,"configuration",lambda:({},pilot.load_template(),b"ref","different"))
    with pytest.raises(HTTPException) as error:ai.prepare(original())
    assert error.value.status_code==409 and not calls
    monkeypatch.setattr(ai,"configuration",lambda:({},pilot.load_template(),b"ref","signature"))
    with pytest.raises(HTTPException) as error:ai.prepare(original(order_fields={}))
    assert error.value.status_code==422 and not calls
    assert ai._SLOTS.acquire(blocking=False)
    ai._SLOTS.release()


def test_configuration_status_never_discloses_private_file_paths_or_keys(monkeypatch):
    monkeypatch.setattr(ai.settings,"carton_mark_ai_enabled",False)
    status=ai.status()
    assert status["available"] is False
    assert set(status)=={"available","template_id","name","model","message"}


def test_ai_request_requires_explicit_consent_but_legacy_payload_remains_valid():
    body=dict(factory_id="huaxing",order_id="o",issue_id="i",excel_asset_id="e",expected_revision=1,layout_id="zumiez-review-v1")
    assert SupplierMarkPdfRequest(**body).engine=="customer_layout"
    with pytest.raises(ValidationError):SupplierMarkPdfRequest(**body,engine="deepseek_zumiez")
    assert SupplierMarkPdfRequest(**body,engine="deepseek_zumiez",ai_consent=True).ai_consent
