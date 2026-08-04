import sys
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


BACKEND_DIR = Path(__file__).resolve().parents[1]

if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))


class FakePage:
    def extract_text(self, **_kwargs):
        return "物料号  数量\n00125  8"

    def extract_tables(self, *, table_settings):
        if table_settings["vertical_strategy"] == "lines":
            return []
        return [[["物料号", "数量"], ["00125", "8"]]]


class FakeDocument:
    pages = [FakePage()]

    def close(self):
        return None


def test_pdf_to_excel_creates_a_new_workbook_and_preserves_leading_zeroes(monkeypatch):
    from app.services import pdf_to_excel as service

    monkeypatch.setattr(service.pdfplumber, "open", lambda _stream: FakeDocument())

    result = service.convert_pdf_to_excel(b"%PDF-fake", "../采购/PO-001.pdf")

    assert result.output_file_name == "PO-001_转换结果.xlsx"
    assert result.page_count == 1
    assert result.table_count == 1
    assert result.text_page_count == 0
    assert result.ocr_page_count == 0

    workbook = load_workbook(BytesIO(result.content))
    assert workbook.sheetnames == ["转换说明", "第1页_表格1"]
    worksheet = workbook["第1页_表格1"]
    assert worksheet["A2"].value == "00125"
    assert worksheet["A2"].number_format == "@"
    assert worksheet.auto_filter.ref == "A1:B2"


def test_pdf_to_excel_uses_ocr_only_when_the_page_has_no_native_content(monkeypatch):
    from app.services import pdf_to_excel as service

    class ScannedPage:
        def extract_text(self, **_kwargs):
            return ""

        def extract_tables(self, *, table_settings):
            return []

    class ScannedDocument(FakeDocument):
        pages = [ScannedPage()]

    monkeypatch.setattr(service.pdfplumber, "open", lambda _stream: ScannedDocument())
    monkeypatch.setattr(service, "_ocr_page", lambda _bytes, _index: [["OCR 字段", "OCR 值"]])

    result = service.convert_pdf_to_excel(b"%PDF-scan", "扫描表.pdf")

    assert result.table_count == 0
    assert result.text_page_count == 0
    assert result.ocr_page_count == 1
    workbook = load_workbook(BytesIO(result.content))
    assert workbook["第1页_OCR"]["A1"].value == "OCR 字段"


def test_pdf_to_excel_preserves_native_document_layout_instead_of_fragmenting_text(monkeypatch):
    from app.services import pdf_to_excel as service

    class LayoutPage:
        width = 600
        height = 800

        def extract_text(self, **_kwargs):
            return "Buyer:  Simba Dickie HK Ltd.\nReference:  SC700149043/500"

        def extract_tables(self, *, table_settings):
            if table_settings["vertical_strategy"] == "lines":
                return []
            return [[
                ["Bu", "yer:", "Si", "mba"],
                ["Re", "fer", "en", "ce"],
                ["SC", "70", "01", "49"],
                ["HK", "Lt", "d.", ""],
                ["Pa", "ge", ":", "1"],
                ["Se", "ll", "er", ":"],
                ["Da", "te", ":", ""],
                ["To", "ta", "l", ""],
            ]]

        def extract_words(self, **_kwargs):
            return [
                {"text": "Buyer:", "x0": 30, "x1": 72, "top": 80, "bottom": 91, "size": 10, "fontname": "Arial"},
                {"text": "Simba", "x0": 98, "x1": 132, "top": 80, "bottom": 91, "size": 10, "fontname": "Arial"},
                {"text": "Dickie", "x0": 136, "x1": 170, "top": 80, "bottom": 91, "size": 10, "fontname": "Arial"},
                {"text": "HK", "x0": 174, "x1": 190, "top": 80, "bottom": 91, "size": 10, "fontname": "Arial"},
                {"text": "Ltd.", "x0": 194, "x1": 220, "top": 80, "bottom": 91, "size": 10, "fontname": "Arial"},
                {"text": "Reference:", "x0": 320, "x1": 380, "top": 140, "bottom": 151, "size": 10, "fontname": "Arial-Bold"},
                {"text": "SC700149043/500", "x0": 388, "x1": 490, "top": 140, "bottom": 151, "size": 10, "fontname": "Arial"},
            ]

    class LayoutDocument(FakeDocument):
        pages = [LayoutPage()]

    monkeypatch.setattr(service.pdfplumber, "open", lambda _stream: LayoutDocument())

    result = service.convert_pdf_to_excel(b"%PDF-layout", "订单.pdf")

    assert result.table_count == 0
    assert result.text_page_count == 1
    assert result.ocr_page_count == 0
    workbook = load_workbook(BytesIO(result.content))
    assert workbook.sheetnames == ["转换说明", "第1页_版式"]
    values = [cell.value for row in workbook["第1页_版式"].iter_rows() for cell in row if cell.value]
    assert "Buyer:" in values
    assert "Simba Dickie HK Ltd." in values
    assert "Reference: SC700149043/500" in values
    assert workbook["第1页_版式"].merged_cells.ranges


def test_pdf_to_excel_replaces_corrupt_native_cjk_mapping_with_ocr_layout(monkeypatch):
    from app.services import pdf_to_excel as service

    class CorruptTextPage:
        width = 600
        height = 800

        def extract_text(self, **_kwargs):
            return "此" * 60 + "箱箱箱箱請請請請" * 8

        def extract_tables(self, *, table_settings):
            return []

    class CorruptTextDocument(FakeDocument):
        pages = [CorruptTextPage()]

    monkeypatch.setattr(service.pdfplumber, "open", lambda _stream: CorruptTextDocument())
    monkeypatch.setattr(
        service,
        "_ocr_page_layout",
        lambda *_args: [service.LayoutBlock("目的地：法國", 100, 160, 260, 175, 10, False)],
    )

    result = service.convert_pdf_to_excel(b"%PDF-corrupt", "繁體訂單.pdf")

    assert result.table_count == 0
    assert result.text_page_count == 0
    assert result.ocr_page_count == 1
    workbook = load_workbook(BytesIO(result.content))
    values = [cell.value for row in workbook["第1页_OCR版式"].iter_rows() for cell in row if cell.value]
    assert "目的地：法國" in values


def test_pdf_to_excel_rejects_documents_without_extractable_content(monkeypatch):
    import pytest
    from app.services import pdf_to_excel as service

    class BlankPage:
        def extract_text(self, **_kwargs):
            return ""

        def extract_tables(self, *, table_settings):
            return []

    class BlankDocument(FakeDocument):
        pages = [BlankPage()]

    monkeypatch.setattr(service.pdfplumber, "open", lambda _stream: BlankDocument())
    monkeypatch.setattr(service, "_ocr_page", lambda _bytes, _index: [])

    with pytest.raises(service.PdfToExcelConversionError, match="未识别到"):
        service.convert_pdf_to_excel(b"%PDF-blank", "空白.pdf")
