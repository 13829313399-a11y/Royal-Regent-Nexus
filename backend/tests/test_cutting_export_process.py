from io import BytesIO
from types import SimpleNamespace
import subprocess
import json
from zipfile import ZipFile
from hashlib import sha256

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from app.services import cutting_export_process as renderer


def report():
    return dict(kind='payments', month='2026-10', token='fixed', captured_at='2026-10-21',
        metadata=[], sheets=[dict(name='精度', headers=['币种', '金额', '货号'],
            rows=[['CNY', '9999999999999999.99', '000001'], ['HKD', '0.00', '=literal']])])


def test_real_process_preserves_values_and_reports_rejected_source():
    actor = SimpleNamespace(id='synthetic', display_name='验证员')
    content, extension = renderer.package(report(), actor, '验证表')
    assert extension == 'xlsx'
    book = load_workbook(BytesIO(content))
    assert list(book['精度'].values)[2:] == [tuple(r) for r in report()['sheets'][0]['rows']]
    assert book['精度']['C4'].data_type == 's'
    value = report()
    value['sheets'][0]['rows'][0][2] = 'x' * 32768
    with pytest.raises(HTTPException) as error:
        renderer.package(value, actor, '验证表')
    assert error.value.status_code == 413


@pytest.mark.parametrize('failure', ['timeout', 'exit', 'truncated', 'start', 'array', 'invalid_error'])
def test_process_failures_never_return_partial_download(monkeypatch, failure):
    def run(command, **kwargs):
        assert kwargs['env']['DATABASE_URL'] == 'sqlite://'
        assert kwargs['timeout'] == 60
        if failure == 'timeout':
            raise subprocess.TimeoutExpired(command, 60)
        if failure == 'start':
            raise OSError('start failed')
        output = b'[]\ncontent' if failure == 'array' else b'{"error":[],"status":"400"}\n' if failure == 'invalid_error' else b'partial'
        return SimpleNamespace(returncode=1 if failure == 'exit' else 0, stdout=output)
    monkeypatch.setattr(renderer.subprocess, 'run', run)
    with pytest.raises(HTTPException) as error:
        renderer.package(report(), SimpleNamespace(id='x', display_name='y'), '验证表')
    assert error.value.status_code == 503


def test_real_process_zip_covers_every_source_row_once():
    value = report()
    # Production MAX_ROWS is 20,000: real multi-volume output without mocking
    # constants in only the parent process.
    expected = [['CNY', str(i), '000001'] for i in range(20001)]
    value['sheets'][0]['rows'] = expected
    content, extension = renderer.package(value, SimpleNamespace(id='x', display_name='y'), '分卷验证')
    assert extension == 'zip'
    with ZipFile(BytesIO(content)) as archive:
        manifest = json.loads(archive.read('manifest.json'))
        collected = []
        for item in manifest['files']:
            raw = archive.read(item['name'])
            assert sha256(raw).hexdigest() == item['sha256']
            book = load_workbook(BytesIO(raw), read_only=True)
            for sheet in book:
                if sheet.title.startswith('精度'):
                    collected.extend([list(row) for row in list(sheet.values)[2:]])
            book.close()
        assert collected == expected
