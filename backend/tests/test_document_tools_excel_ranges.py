import hashlib

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Border, Font, PatternFill, Side
import pytest

from app.services.document_tools import office_engine as office
from app.services.document_tools.document_ir import Cancelled, ToolError


def extract(path, options=None, cancelled=lambda: False):
    return office._extract_xlsx(path, options or {}, lambda *_: None, cancelled)


def sample(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "订单"
    sheet.append(["订单号", "数量"])
    sheet.append(["00001234567890123456789", 12])
    # Excel remembers formatted cells far beyond the real table.
    sheet["Z10000"].number_format = "0.00"
    sheet["Z10000"].font = Font(name="Arial", size=11)
    return book, sheet, tmp_path / "format-only-tail.xlsx"


def test_automatic_range_excludes_invisible_format_only_tail(tmp_path):
    book, sheet, path = sample(tmp_path)
    book.save(path)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    ir = extract(path)
    table = ir.tables[0]
    assert (table.row_count, table.column_count) == (2, 2)
    assert table.source.cell == "A1:B2"
    assert any(c.value == "00001234567890123456789" for c in table.cells)
    assert any(i.code == "EXCEL_UNUSED_RANGE_TRIMMED" for i in ir.issues)
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_visible_blank_template_cells_and_merges_are_preserved(tmp_path):
    book, sheet, path = sample(tmp_path)
    sheet["D6"].border = Border(bottom=Side(style="thin"))
    sheet["F8"].fill = PatternFill("solid", fgColor="FFFF00")
    sheet.merge_cells("B9:D10")
    book.save(path)
    table = extract(path).tables[0]
    assert table.source.cell == "A1:F10"
    assert (table.row_count, table.column_count) == (10, 6)
    assert any(c.source.cell == "B9" and c.rowspan == 2 and c.colspan == 3 for c in table.cells)


@pytest.mark.parametrize("kind", ["value", "formula", "border", "fill", "merge", "conditional", "row", "column"])
def test_real_distant_content_is_never_silently_dropped(tmp_path, kind):
    book, sheet, path = sample(tmp_path)
    if kind == "value":
        sheet["Z10000"] = 0
    elif kind == "formula":
        sheet["Z10000"] = '=IF(A1="","",0)'
    elif kind == "border":
        sheet["Z10000"].border = Border(bottom=Side(style="thin"))
    elif kind == "fill":
        sheet["Z10000"].fill = PatternFill("solid", fgColor="FFFFFF")
    elif kind == "merge":
        sheet.merge_cells("Y9999:Z10000")
    elif kind == "conditional":
        sheet.conditional_formatting.add("Z10000", FormulaRule(formula=["1=1"],
                                           fill=PatternFill("solid", fgColor="FFFF00")))
    elif kind == "row":
        sheet["Z1"] = "Last data column"
        sheet.row_dimensions[10000].fill = PatternFill("solid", fgColor="FFFF00")
    else:
        sheet["A10000"] = "Last data row"
        sheet.column_dimensions["Z"].fill = PatternFill("solid", fgColor="FFFF00")
    book.save(path)
    with pytest.raises(ToolError) as caught:
        extract(path)
    assert caught.value.code == "OFFICE_RANGE_TOO_LARGE"
    assert "订单" in str(caught.value)
    assert "A1:Z10000" in str(caught.value)


@pytest.mark.parametrize("explicit_print_area", [False, True])
def test_explicit_ranges_keep_their_original_boundaries(tmp_path, explicit_print_area):
    book, sheet, path = sample(tmp_path)
    options = {"range": "A1:Z10000"}
    if explicit_print_area:
        sheet.print_area = "A1:Z10000"
        options = {}
    book.save(path)
    with pytest.raises(ToolError) as caught:
        extract(path, options)
    assert caught.value.code == "OFFICE_RANGE_TOO_LARGE"
    ir = extract(path, {"range": "A1:B2"})
    assert ir.tables[0].source.cell == "A1:B2"
    assert not any(i.code == "EXCEL_UNUSED_RANGE_TRIMMED" for i in ir.issues)


def test_small_original_template_range_is_not_changed(tmp_path):
    book = Workbook()
    book.active["A1"] = "Header"
    book.active["D10"].number_format = "0.00"
    path = tmp_path / "small.xlsx"
    book.save(path)
    ir = extract(path)
    assert ir.tables[0].source.cell == "A1:D10"
    assert not any(i.code == "EXCEL_UNUSED_RANGE_TRIMMED" for i in ir.issues)


@pytest.mark.parametrize("axis", ["row", "column"])
def test_background_band_does_not_expand_into_perpendicular_ghost_formatting(tmp_path, axis):
    book, sheet, path = sample(tmp_path)
    dimension = sheet.row_dimensions[1] if axis == "row" else sheet.column_dimensions["A"]
    dimension.fill = PatternFill("solid", fgColor="FFFF00")
    book.save(path)
    assert extract(path).tables[0].source.cell == "A1:B2"


def test_cancellation_still_stops_extraction(tmp_path):
    book, sheet, path = sample(tmp_path)
    book.save(path)
    with pytest.raises(Cancelled):
        extract(path, cancelled=lambda: True)


def test_calc_format_expansion_does_not_change_accepted_source_range(tmp_path, monkeypatch):
    book, sheet, path = sample(tmp_path)
    book.save(path)
    seen_ranges = []

    def render(source, output, options, work_dir, cancelled):
        if output.suffix == ".pdf":
            raise ToolError("OFFICE_UNAVAILABLE", "Preview unavailable in this native test")
        seen_ranges.append(options["_native_ranges"])
        # Simulate Calc exporting formatting as visible empty cells far away.
        sheet["Z10000"].border = Border(bottom=Side(style="thin"))
        book.save(output)
        return {}

    monkeypatch.setattr(office, "office_executable", lambda: "test-office")
    monkeypatch.setattr(office, "render_office", render)
    result = office.convert_office(path, "excel_to_word", {}, tmp_path / "result", lambda *_: None, lambda: False)
    assert seen_ranges == [{"订单": ["A1:B2"]}]
    assert result.ir.tables[0].source.cell == "A1:B2"
    assert not any(i.code == "OFFICE_RANGE_TOO_LARGE" for i in result.ir.issues)
    assert (tmp_path / "result/output.docx").is_file()
