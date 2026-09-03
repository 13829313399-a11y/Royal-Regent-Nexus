from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image

from app.services.internal_quote_excel import _build_summary_sheet
from test_internal_quote_excel_formula_links import _section, _find_row


def test_department_blocks_prices_images_and_formulas_are_component_scoped():
    components = [{"id": "a", "name": "电话", "cost_hkd": "10", "markup": "1.16"},
                  {"id": "b", "name": "镜子", "cost_hkd": "7", "markup": "1.25"}]
    entries = [{"section": code, "kind": "manual", "label": label, "amount_hkd": cost, "pricing_component_id": component}
               for component, code, label, cost in [("a", "engineering", "电话壳", 3), ("a", "electronic", "电话IC", 2),
                                                    ("a", "sewing", "电话布套", 4), ("a", "hair", "电话车发", 1),
                                                    ("b", "assembly", "镜子装配", 5), ("b", "electronic", "镜子LED", 2)]]
    attachments = []
    for component, color in [("a", "red"), ("b", "blue"), ("legacy", "green")]:
        buffer = BytesIO()
        Image.new("RGB", (300, 500), color).save(buffer, format="PNG")
        attachments.append(SimpleNamespace(id=component, department="product-image" if component == "legacy" else f"component-image:{component}",
                                           content_type="image/png", content=buffer.getvalue(), uploaded_at="2026-09-02"))
    summary = {"t1": [], "t2": [], "t3": [], "t4": [], "shipping_pricing": {
        "pricing_mode": "component", "pricing_groups": components, "pricing_entries": entries,
        "markup": 1.2, "misc_ratio": .03, "global_pricing": {"entries": [], "cost_hkd": 0, "markup": 1.3}, "rows": [],
    }}
    workbook = Workbook()
    _build_summary_sheet(workbook, SimpleNamespace(product_name="分项验证", quote_no="JP-CHECK", customer="JustPlay", region_code="mainland", remark=""),
                         [_section("sales", payload={"pricing_mode": "component", "pricing_components": components})],
                         {"fx": {"rmb_hkd": .85, "hkd_usd": 7.8}}, summary, {}, attachments)
    sheet = workbook["报价明细"]
    titles = ["电话明细", "电话 · 电子", "电话 · 车缝", "电话 · 车发", "镜子明细", "镜子 · 电子"]
    starts = [_find_row(sheet, 3, title) for title in titles]
    assert starts == sorted(starts)
    assert not any(sheet.cell(row, 3).value in {"镜子 · 车缝", "镜子 · 车发"} for row in range(1, sheet.max_row + 1))
    expected_costs = [3, 2, 4, 1, 5, 2]
    usd_refs = []
    total_usd = 0
    for index, start in enumerate(starts):
        end = next(row for row in range(start + 1, sheet.max_row + 1) if sheet.cell(row, 3).value == "成本金额：")
        assert sheet.cell(end, 4).value == f"=SUM(D{start+1}:D{end-1})"
        if " · 电子" in titles[index]:
            assert end == start + 6
            assert sheet.cell(start + 4, 4).value == f"=H{start}-SUM(D{start+1}:D{start+3})"
            cost = float(sheet.cell(start, 8).value[1:]) + sheet.cell(start + 5, 4).value
        else:
            cost = sum(float(sheet.cell(row, 4).value) for row in range(start + 1, end))
        assert cost == expected_costs[index]
        markup = sheet.cell(end + 1, 4).value
        assert markup == (1.16 if index < 4 else 1.25)
        assert sheet.cell(end + 2, 4).value == "=1-$Q$6"
        assert sheet.cell(end + 3, 4).value == f"=D{end}*D{end+1}/D{end+2}"
        assert sheet.cell(end + 5, 4).value == f"=D{end+3}/D{end+4}"
        usd_refs.append(f"$D${end+5}")
        total_usd += cost * markup / .97 / 7.8
    assert total_usd == pytest.approx((10 * 1.16 + 7 * 1.25) / .97 / 7.8)
    combined = [cell.value for row in sheet for cell in row if cell.data_type == "f" and all(ref in cell.value for ref in usd_refs)]
    assert combined and all(combined[0].count(ref) == 1 for ref in usd_refs)
    assert len(sheet._images) == 2
    assert [i.anchor._from.row for i in sheet._images] == [_find_row(sheet, 1, "电话-明细") - 1, _find_row(sheet, 1, "镜子-明细") - 1]
    assert all(i.anchor._from.col == 13 for i in sheet._images)
    assert all(i.height / i.width == pytest.approx(5/3, abs=.02) for i in sheet._images)
    for row in sheet:
        for cell in row:
            if cell.data_type == "f":
                assert "#REF!" not in cell.value
    output = BytesIO()
    workbook.save(output)
    reopened = load_workbook(BytesIO(output.getvalue()))
    assert len(reopened["报价明细"]._images) == 2
    assert reopened["报价明细"].cell(starts[0], 1).border.top.style == "medium"
    reopened.close()
    workbook.close()
