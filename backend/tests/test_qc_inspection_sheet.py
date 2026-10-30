"""Real SQLite / HTTP / XLSX acceptance. No persistent user data is touched."""
import base64
from copy import deepcopy
from dataclasses import replace
from io import BytesIO
import os
from pathlib import Path
from uuid import uuid4
import zipfile

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
import httpx
from openpyxl import load_workbook
from PIL import Image
import pytest

from test_customer_order_ledger import engine, db, ingest, preview
from test_qc_workflow import qc_user, ACTOR, create as create_workflow
from app.api import qc_department as gateway
from app.db import get_db
from app.services.auth import get_current_user
from app.services import qc_inspection_sheet as sheet, qc_workflow as wf


@pytest.fixture(autouse=True)
def isolated(monkeypatch,tmp_path):
    monkeypatch.setattr(sheet.settings,'qc_department_data_dir',str(tmp_path/'companies'))


def source(identity='order-a',po='000001'):
    return {'id':identity,'customer':'客户甲','poNo':po,'customerPo':'001019','productNo':'00042',
            'productName':'同款产品','dateFrom':'2026-10-29','dateTo':'2026-10-29','quantity':'24',
            'unitsPerCarton':'6','cartonCount':'4','production':'装配','sourceStatus':'active','remark':''}


def pushed_candidates(sources):
    tasks=[{'id':'task-'+s['id'],'sourceId':s['id'],'source':s,'batch':'首验','execution':{},'attachments':[]} for s in sources]
    return sheet.candidates('huaxing',[],tasks,sources)


def old_source_candidates(sources):
    # Reproduce pre-fix cached rows only for migration compatibility coverage.
    return [('source:'+s['id'],{**sheet.base_fields(s),'sourceId':s['id']},'',s.get('remark',''),[]) for s in sources]


def seed():
    rows=pushed_candidates([source('a'),source('b','000002'),source('c','000003')])
    return sheet.sync('huaxing',rows)['items']


def save(row,result='',remark=''):
    return sheet.save('huaxing',row['id'],{'revision':row['revision'],'result':result,'remark':remark},ACTOR)


def photo(row,color='red',key=None,fmt='PNG'):
    output=BytesIO();Image.new('RGB',(160,90),color).save(output,format=fmt)
    body={'revision':row['revision'],'name':'现场照片.'+fmt.lower(),'content':base64.b64encode(output.getvalue()).decode(),'requestKey':key or uuid4().hex}
    return sheet.add_photo('huaxing',row['id'],body,ACTOR),body


def test_stable_identity_automatic_source_mapping_duplicates_and_preserved_input():
    a,b,c=seed();assert all(r['result']=='' for r in (a,b,c))
    assert len({r['id'] for r in (a,b,c)})==3 and a['productNo']==b['productNo']
    assert a['country']=='' and a['agency']=='' and a['poNo']=='000001'
    a=save(a,'PASS','中文 English 0123 !@\n长备注'*200)
    a=save(a,'FAIL',a['remark']+'修改')
    rows=sheet.sync('huaxing',pushed_candidates([source('c','000003'),source('b','000002'),source('a')]))['items']
    assert next(r for r in rows if r['id']==a['id'])['result']=='FAIL'
    assert next(r for r in rows if r['id']==b['id'])['result']==''
    changed=source('a');changed['quantity']='48'
    rows=sheet.sync('huaxing',pushed_candidates([changed]))['items']
    assert rows[0]['quantity']=='48' and rows[0]['remark']==a['remark']
    retained=sheet.sync('huaxing',[])['items'];assert len(retained)==1 and retained[0]['result']=='FAIL'
    assert retained[0]['sourceWarning']
    with sheet.store() as conn:assert conn.execute('SELECT count(*) FROM changes').fetchone()[0]==2


def test_optimistic_conflicts_invalid_results_and_factory_isolation():
    a,b,c=seed();saved=save(a,'PASS')
    with pytest.raises(HTTPException) as error:save(a,'FAIL')
    assert error.value.status_code==409
    for value in ['Pending','Approved','REJ','COND',None]:
        with pytest.raises(HTTPException) as error:save(saved,value)
        assert error.value.status_code==422
    with pytest.raises(HTTPException) as error:sheet.save('huadeng',a['id'],{'revision':saved['revision'],'result':'FAIL','remark':''},ACTOR)
    assert error.value.status_code==404
    assert save(saved,'')['result']==''


def test_photos_persist_idempotency_delete_validation_and_task_scope():
    a,b,c=seed();a,body=photo(a);one=a['photos'][0]['id'];a,second=photo(a,'blue')
    assert len(a['photos'])==2
    assert len(sheet.add_photo('huaxing',a['id'],body,ACTOR)['photos'])==2
    content=sheet.read_photo('huaxing',a['id'],one);assert Image.open(BytesIO(content)).format=='JPEG'
    for factory,entry in [('huadeng',a['id']),('huaxing',b['id'])]:
        with pytest.raises(HTTPException):sheet.read_photo(factory,entry,one)
    a=sheet.remove_photo('huaxing',a['id'],{'revision':a['revision'],'photoId':one},ACTOR)
    assert len(a['photos'])==1
    with pytest.raises(HTTPException):sheet.read_photo('huaxing',a['id'],one)
    with sheet.store() as conn:assert conn.execute('SELECT removed FROM photos WHERE id=?',(one,)).fetchone()[0]==1
    body.update(revision=a['revision'],requestKey=uuid4().hex,content=base64.b64encode(b'not a photo').decode())
    with pytest.raises(HTTPException) as error:sheet.add_photo('huaxing',a['id'],body,ACTOR)
    assert error.value.status_code==422
    with pytest.raises(HTTPException) as error:sheet.normalize_photo(b'x'*(10*1024*1024+1))
    assert error.value.status_code==413


def test_mobile_heif_and_webp_are_persistent_jpeg():
    from pillow_heif import from_pillow
    output=BytesIO();from_pillow(Image.new('RGB',(32,16),'green')).save(output)
    assert Image.open(BytesIO(sheet.normalize_photo(output.getvalue()))).format=='JPEG'
    output=BytesIO();Image.new('RGB',(32,16),'blue').save(output,format='WEBP')
    assert Image.open(BytesIO(sheet.normalize_photo(output.getvalue()))).format=='JPEG'


def test_export_three_states_actual_embedded_images_text_ids_and_revision_gate():
    a,b,c=seed();a=save(a,'PASS','合格 中文 English');a,_=photo(a)
    b=save(b,'FAIL','=备注原文\n外壳划痕');b,_=photo(b,'blue');b,_=photo(b,'green')
    body={'entries':[{'id':r['id'],'revision':r['revision']} for r in (a,b,c)]}
    content=sheet.export_xlsx('huaxing',body)
    book=load_workbook(BytesIO(content));ws=book['验货明细']
    assert [c.value for c in ws[1]]==[label for _,label in sheet.COLUMNS]
    assert [c.value for c in ws[1]][-3:]==['验货结果','备注','照片']
    assert ws.max_row==4 and ws['N2'].value=='PASS' and ws['N3'].value=='FAIL' and ws['N4'].value in ('',None)
    assert ws['O2'].value==a['remark'] and ws['O3'].value==b['remark'] and ws['O3'].data_type=='s'
    assert ws['C2'].value=='000001' and ws['E3'].value=='00042' and ws['L2'].value=='001019'
    assert ws['A2'].value.date().isoformat()=='2026-10-29' and ws['G2'].value==24
    assert len(ws._images)==3 and [i.anchor._from.row for i in ws._images]==[1,2,2]
    assert all(i.anchor._from.col==15 for i in ws._images)
    assert ws._images[1].anchor._from.colOff != ws._images[2].anchor._from.colOff
    with zipfile.ZipFile(BytesIO(content)) as archive:
        assert len([n for n in archive.namelist() if n.startswith('xl/media/')])==3
        assert archive.testzip() is None
    if os.environ.get('QC_INLINE_EVIDENCE'):
        Path(os.environ['QC_INLINE_EVIDENCE']).write_bytes(content)
    save(c,'PASS')
    with pytest.raises(HTTPException) as error:sheet.export_xlsx('huaxing',body)
    assert error.value.status_code==409


def test_all_photos_export_even_more_than_eleven():
    a,*_=seed()
    for n in range(12):a,_=photo(a)
    book=load_workbook(BytesIO(sheet.export_xlsx('huaxing',{'entries':[{'id':a['id'],'revision':a['revision']}]})))
    assert len(book['照片-1']._images)==12
    assert book['验货明细']['P2'].hyperlink.target=="#'照片-1'!A1"
    assert all(book['照片-1'].cell(r,3).value==a['id'] for r in range(2,14))


def test_existing_workflow_and_legacy_snapshots_preserved(db):
    task=create_workflow(db);before=deepcopy(wf.detail(db,'huaxing',task['id']))
    entries=sheet.sync('huaxing',sheet.candidates('huaxing',[],[task],wf.sources(db,'huaxing')['items']))['items']
    assert len(entries)==1
    saved=save(entries[0],'FAIL','新行内备注');photo(saved)
    assert wf.detail(db,'huaxing',task['id'])==before
    assert wf.projections(db,'huaxing')==[]
    legacy={'id':7,'client':'甲','orderNo':'001','productNo':'012','productName':'零件','result':'COND','remark':'历史条件'}
    before=deepcopy(legacy)
    entries=sheet.sync('huaxing',sheet.candidates('huaxing',[legacy],[],[]))['items']
    old=next(r for r in entries if r['sourceKey']=='legacy:7')
    assert old['result']=='' and old['historicalResult']=='COND' and legacy==before


def test_http_actual_save_read_photo_export_permissions_and_source_unchanged(db,monkeypatch):
    from app.services import auth
    monkeypatch.setattr(auth,'add_auth_audit',lambda *a,**k:None)
    async def upstream(*args):return httpx.Response(200,json={'records':[]})
    monkeypatch.setattr(gateway,'upstream',upstream)
    line=ingest(db);original=deepcopy(line.data)
    app=FastAPI();app.include_router(gateway.router);current=[qc_user()]
    app.dependency_overrides[get_current_user]=lambda:current[0];app.dependency_overrides[get_db]=lambda:db
    with TestClient(app) as client:
        base=client.get('/api/qc-department/huaxing/api/health').json()['frame_url']
        assert client.get(base+'api/inspection-sheet/entries').json()['items']==[]
        response=client.post(base+'api/workflow/tasks',json={'requestKey':uuid4().hex,'items':[{'sourceId':line.id,'sourceRevision':line.revision,'method':'internal','plannedDate':'2026-10-29','plannedQuantity':str(line.quantity),'batch':'首验'}]})
        assert response.status_code==200,response.text
        schedule=client.get(base+'api/schedule-source').json()
        row=client.get(base+'api/inspection-sheet/entries').json()['items'][0]
        path=base+'api/inspection-sheet/entries/'+row['id']
        payload={'revision':row['revision'],'result':'PASS','remark':'现场记录'}
        response=client.post(path+'/save',json=payload);assert response.status_code==200;row=response.json()
        assert client.post(path+'/save',json=payload).status_code==409
        assert client.post(path+'/save',json=payload,headers={'Sec-Fetch-Site':'cross-site'}).status_code==403
        output=BytesIO();Image.new('RGB',(10,10),'red').save(output,format='PNG')
        response=client.post(path+'/photos',json={'revision':row['revision'],'name':'照片.png','content':base64.b64encode(output.getvalue()).decode(),'requestKey':uuid4().hex})
        assert response.status_code==200;row=response.json()
        assert client.get(path+'/photos/'+row['photos'][0]['id']).headers['content-type']=='image/jpeg'
        ticket=client.post(base+'api/inspection-sheet/export',json={'entries':[{'id':row['id'],'revision':row['revision']}]}).json()
        download=client.get(base+ticket['url']);assert download.status_code==200
        assert len(load_workbook(BytesIO(download.content)).active._images)==1
        reloaded=client.get(base+'api/inspection-sheet/entries').json()['items'][0]
        assert reloaded['result']=='PASS' and reloaded['remark']=='现场记录' and len(reloaded['photos'])==1
        assert client.get(base+'api/schedule-source').json()==schedule and line.data==original
        assert client.get('/api/qc-department/huadeng/api/health').status_code==403
        current[0]=qc_user(True)
        assert client.get(path+'/photos/'+row['photos'][0]['id']).status_code==409
        newbase=client.get('/api/qc-department/huaxing/api/health').json()['frame_url']
        assert client.get(newbase+ticket['url']).status_code==404


def test_unpushed_orders_and_old_caches_stay_hidden_until_confirmed(db):
    line=ingest(db);sources=wf.sources(db,'huaxing')['items']
    assert sheet.sync('huaxing',sheet.candidates('huaxing',[],[],sources))['items']==[]
    old=sheet.sync('huaxing',old_source_candidates(sources))['items'][0]
    old=save(old,'FAIL','保留旧版已填内容')
    assert sheet.sync('huaxing',sheet.candidates('huaxing',[],[],sources))['items']==[]
    task=create_workflow(db,line=line)
    rows=sheet.sync('huaxing',sheet.candidates('huaxing',[],[task],sources))['items']
    assert len(rows)==1 and rows[0]['id']==old['id'] and rows[0]['result']=='FAIL' and rows[0]['remark']=='保留旧版已填内容'
    assert len(sheet.sync('huaxing',sheet.candidates('huaxing',[],[task],sources))['items'])==1


def test_pending_migrates_old_check_without_losing_photos_changes_or_ids():
    import sqlite3
    a=save(seed()[0],'FAIL','existing note');a,_=photo(a)
    path=Path(sheet.settings.qc_department_data_dir)/'_inspection_sheet/inspections.sqlite3'
    with sqlite3.connect(path) as conn:
        conn.execute('PRAGMA foreign_keys=OFF')
        conn.execute(sheet.SCHEMA.split(';')[0].replace('IF NOT EXISTS entries','entries_old').replace(",'PENDING'",''))
        conn.execute('INSERT INTO entries_old SELECT * FROM entries');conn.execute('DROP TABLE entries');conn.execute('ALTER TABLE entries_old RENAME TO entries')
        before={t:list(conn.execute('SELECT * FROM '+t)) for t in ('entries','photos','changes','exports')}
    with sheet.store() as conn:
        for table,rows in before.items():assert [tuple(r) for r in conn.execute('SELECT * FROM '+table)]==rows
        assert not conn.execute('PRAGMA foreign_key_check').fetchall()
    saved=save(a,'PENDING','待定，等客户确认')
    assert saved['result']==saved['cumulativeResult']=='PENDING'
    assert len(saved['photos'])==1
    with sheet.store() as conn:assert sheet.get(conn,'huaxing',a['id'])['result']=='PENDING'
