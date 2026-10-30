"""Mixed receipt cartons preserve quantities and count physical boxes exactly once."""
from decimal import Decimal
from io import BytesIO
import json

import openpyxl
import pytest

from test_indonesia_materials import client, command, material, product, order, workspace, BASE
from test_indonesia_warehouse_import import source, preview, choose, confirm
from test_indonesia_receipt_template import download, fill
from app.services.indonesia_receipt_template import HEADERS, HEADER_ROW, SHEET
from app.services.indonesia_material_excel import SHEET_NAME
from app.services.indonesia_warehouse_packing import packing_groups

PACKING = dict(package_label='M01', cartons='2', length_cm='27', width_cm='14.5', height_cm='27.3', carton_weight_kg='.02')


def test_large_incomplete_group_keeps_review_metadata_bounded():
    rows = [dict(index=index, source=source(package_label='M01', receipt_line=str(index + 1))) for index in range(2000)]
    group = packing_groups(rows)[0]
    assert len(group['issues']) == 4 and '1997' in group['issues'][-1]
    metadata = {key: value for key, value in group.items() if key not in {'indexes', 'package'}}
    assert len(json.dumps(metadata, ensure_ascii=False).encode()) < 1200


def setup_mixed(client):
    first = material(client)
    second = material(client, name='合成拉杆', net_weight_kg='.002')  # Shared code, distinct official name.
    product(client, first, materials=[dict(material_id=item['id'], usage='1') for item in (first, second)])
    contract = order(client, quantity='200')
    return first, second, contract


def mixed_rows(**changes):
    return [source(**PACKING), source(**(PACKING | dict(material_name='合成拉杆', receipt_line='2', purchase_line='0002',
                   normal_quantity='50', spare_quantity='5') | changes))]


def shipment_for(client, imported):
    return next(row for row in workspace(client)['records'] if row['id'] == imported['data']['shipment_ids'][-1])


def test_same_box_two_materials_are_one_group_through_confirm_and_excel(client):
    setup_mixed(client)
    imported = preview(client, *mixed_rows(cartons='2.0', length_cm='27.00'))
    groups = [row['packing_group'] for row in imported['data']['rows']]
    assert groups[0]['id'] == groups[1]['id'] and groups[0]['row_count'] == 2 and not groups[0]['issues']
    saved = confirm(client, imported, [choose(imported, i) for i in range(2)])
    shipment = shipment_for(client, saved)
    data = shipment['data']
    assert data['packing_pending'] is False and not data['packing_notes']
    assert len(data['packages']) == 1 and len(data['packages'][0]['items']) == 2
    assert data['packages'][0]['cartons'] == 2
    done = command(client, 'shipment_confirm', {}, shipment)
    totals = done['data']['packing_totals']
    assert totals['cartons'] == 2 and Decimal(totals['cbm']) == Decimal('.02')
    assert sum(Decimal(line['net_weight_kg']) for line in totals['lines'].values()) == Decimal('.21')
    assert sum(Decimal(line['gross_weight_kg']) for line in totals['lines'].values()) == Decimal('.25')
    response = client.get(BASE + f"/shipments/{done['id']}/export", params={'factory_id': 'huaxing'})
    assert response.status_code == 200
    book = openpyxl.load_workbook(BytesIO(response.content), data_only=True)
    detail = book[SHEET_NAME]
    names = [cell.value for row in detail for cell in row]
    assert '合成螺丝' in names and '合成拉杆' in names
    assert {'AT4:AT5', 'S4:S5'} <= {str(area) for area in detail.merged_cells.ranges}
    assert detail['AT4'].value == 2 and detail['AT5'].value is None
    assert detail['S4'].value == pytest.approx(.02) and detail['S5'].value is None


@pytest.mark.parametrize('changes,reason', [
    ({'cartons': '3'}, '箱数不一致'), ({'length_cm': '28'}, '箱长不一致'),
    ({'carton_weight_kg': '.03'}, '每箱纸箱重量不一致'), ({'width_cm': ''}, '缺少箱宽'),
    ({'cartons': '1.5'}, '装箱数值无效'), ({'height_cm': '0'}, '装箱数值无效'),
    ({'actual_date': '2026-10-23'}, '实收日期不一致'),
])
def test_inconsistent_or_incomplete_box_is_reviewable_and_never_auto_packed(client, changes, reason):
    setup_mixed(client)
    imported = preview(client, *mixed_rows(**changes))
    assert reason in '；'.join(imported['data']['rows'][0]['packing_group']['issues'])
    saved = confirm(client, imported, [choose(imported, i) for i in range(2)])
    shipment = shipment_for(client, saved)
    assert shipment['data']['packages'] == [] and shipment['data']['packing_pending']
    assert reason in '；'.join(shipment['data']['packing_notes'])
    payload = {key: shipment['data'][key] for key in ('customer', 'recipient', 'shipment_date', 'lines', 'packages', 'remarks')}
    saved_draft = command(client, 'shipment_save', payload, shipment)
    assert saved_draft['data']['packing_notes'] == shipment['data']['packing_notes']
    reopened = next(row for row in workspace(client)['records'] if row['id'] == shipment['id'])
    assert reopened['data']['packing_notes'] == shipment['data']['packing_notes']
    command(client, 'shipment_confirm', {}, reopened, status=422)


@pytest.mark.parametrize('changes', [{'warehouse': '半成品仓'}, {'receipt_no': 'WH-002'}])
def test_same_label_in_different_warehouse_or_receipt_does_not_merge(client, changes):
    setup_mixed(client)
    imported = preview(client, *mixed_rows(**changes))
    saved = confirm(client, imported, [choose(imported, i) for i in range(2)])
    shipment = shipment_for(client, saved)
    assert len(shipment['data']['packages']) == 2
    assert len({p['label'] for p in shipment['data']['packages']}) == 2
    assert not shipment['data']['packing_pending']
    done = command(client, 'shipment_confirm', {}, shipment)
    assert done['data']['packing_totals']['cartons'] == 4


def test_partial_selection_and_later_import_never_count_incomplete_mixed_box(client):
    setup_mixed(client)
    imported = preview(client, *mixed_rows())
    first = confirm(client, imported)
    assert shipment_for(client, first)['data']['packages'] == []
    refreshed = next(row for row in workspace(client)['records'] if row['id'] == imported['id'])
    second = confirm(client, refreshed, [choose(refreshed, 1)])
    data = shipment_for(client, second)['data']
    assert data['packing_pending'] and data['packages'] == []
    assert '同箱已有导入记录' in '；'.join(data['packing_notes'])


def test_later_file_for_already_received_box_requires_review(client):
    setup_mixed(client)
    first_row, second_row = mixed_rows()
    confirm(client, preview(client, first_row))
    imported = preview(client, second_row)
    assert '同箱已有导入记录' in '；'.join(imported['data']['rows'][0]['packing_group']['issues'])
    saved = confirm(client, imported)
    assert shipment_for(client, saved)['data']['packages'] == []


def test_order_template_mixed_box_roundtrip_keeps_distinct_names_and_quantities(client):
    _, _, contract = setup_mixed(client)
    template = download(client, contract).content
    book = openpyxl.load_workbook(BytesIO(template))
    assert '混箱' in book[SHEET]['A4'].value
    assert 'A物料100件' in book[SHEET].cell(HEADER_ROW, HEADERS.index('箱号') + 1).comment.text
    for position, row in enumerate(mixed_rows(), 6):
        template = fill(template, {'收料单号': 'WH-001', '行号': row['receipt_line'], '实收日期': row['actual_date'],
            '正常实收数量': row['normal_quantity'], '备品数量': row['spare_quantity'], '箱号': 'M01',
            '箱数': '2', '长cm': '27', '宽cm': '14.5', '高cm': '27.3', '每箱纸箱重量kg': '.02'}, row=position)
    response = client.post(BASE + '/warehouse-imports/preview', data=dict(factory_id='huaxing', customer='合成客户', recipient='RRI', warehouse=''),
                           files={'file': ('混箱回传.xlsx', template, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code == 200, response.text
    imported = response.json()['record']
    selections = [choose(imported, i, purchase={key: row[key] for key in ('purchase_no', 'purchase_line', 'purchase_date', 'purchase_quantity', 'supplier', 'price', 'currency')})
                  for i, row in enumerate(mixed_rows())]
    posted = confirm(client, imported, selections)
    data = shipment_for(client, posted)['data']
    assert len(data['packages']) == 1 and len(data['packages'][0]['items']) == 2
    assert [(item['normal_quantity'], item['spare_quantity']) for item in data['packages'][0]['items']] == [('90', '10'), ('50', '5')]
