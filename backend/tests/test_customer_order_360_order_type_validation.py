from __future__ import annotations

from copy import deepcopy
from datetime import date
from hashlib import sha256
from io import BytesIO

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.customer_order_360_order_types import (
    ORDER_TYPES, SHEETS, OrderTypeValidationError, ensure_360_order_type_dropdowns,
)
from customer_order_360_private_samples import require_private_manifest


def _workbook():
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name in SHEETS:
        sheet = workbook.create_sheet(name)
        sheet["B3"] = "订单类型"
        sheet["A12"] = "取消单"
        sheet["B4"] = "备料单"
        sheet["B5"] = '=IF(ITEM表!B4="","",ITEM表!B4)' if name != "ITEM表" else "DI Order"
        sheet["B13"] = "人工取消类型"
        sheet["E4"] = "RL-111-1"
        sheet["H4"] = "12345"
    return workbook


def _add(sheet, formula='"正单,加单,补单,改单"', cells="B4:B11", **kwargs):
    validation = DataValidation(type="list", formula1=formula, **kwargs)
    validation.add(cells)
    sheet.add_data_validation(validation)
    return validation


def _apply(workbook):
    return ensure_360_order_type_dropdowns(workbook, factory_id="huakang-a", customer_code="360")


def _rules(sheet, cell):
    return [v for v in sheet.data_validations.dataValidation if cell in v.sqref]


def _snapshot(workbook):
    return {s.title: {(r, c): (cell.value, cell.data_type, cell.style_id)
                     for (r, c), cell in s._cells.items()} for s in workbook}


def _validation_snapshot(workbook):
    return [(s.title, [(str(v.sqref), v.formula1, v.type, v.showDropDown)
                      for v in s.data_validations.dataValidation]) for s in workbook]


def test_three_sheets_include_required_types_preserve_extra_options_and_cell_formulas():
    workbook = _workbook()
    _add(workbook["ITEM表"])
    for name in SHEETS[1:]:
        _add(workbook[name], '"正式PO,备料"', promptTitle="原提示", showErrorMessage=True)
    unrelated = _add(workbook["ITEM表"], '"是,否"', "V4:V20")
    before = _snapshot(workbook)
    result = _apply(workbook)
    assert result["applied"]
    assert _snapshot(workbook) == before
    assert unrelated.formula1 == '"是,否"' and str(unrelated.sqref) == "V4:V20"
    for name in SHEETS:
        rule, = _rules(workbook[name], "B4")
        assert all(option in rule.formula1 for option in ORDER_TYPES)
        assert rule.showDropDown is False
        assert not _rules(workbook[name], "B3")
        assert not _rules(workbook[name], "B12")
        assert not _rules(workbook[name], "B13")
    summary, = _rules(workbook["接单表"], "B4")
    assert summary.formula1 == '"正式PO,备料,正单,加单,补单,改单,备料单"'
    assert summary.promptTitle == "原提示" and summary.showErrorMessage


def test_missing_validation_and_new_rows_receive_complete_dropdown():
    workbook = _workbook()
    for sheet in workbook:
        _add(sheet, cells="B4:B8")
        sheet["B11"] = "备料单"
    _apply(workbook)
    for sheet in workbook:
        for row in range(4, 12):
            rule, = _rules(sheet, f"B{row}")
            assert rule.formula1 == '"正单,加单,补单,改单,备料单"'
    missing = _workbook()
    _apply(missing)
    for sheet in missing:
        assert _rules(sheet, "B11")[0].formula1 == '"正单,加单,补单,改单,备料单"'


def test_merged_and_group_headings_are_excluded_without_rewriting_original_rules():
    workbook = _workbook()
    for sheet in workbook:
        sheet.merge_cells("A7:D7")
        sheet["A7"] = "已完成（按验货日期）"
        sheet["A9"] = "未完成"
        sheet["B10"] = "订单类型"
    # A rule shared with another column and the historical/cancelled area must
    # retain its original semantics everywhere outside editable type cells.
    original = _add(workbook["ITEM表"], cells="A1:C15")
    _apply(workbook)
    sheet = workbook["ITEM表"]
    assert original.formula1 == '"正单,加单,补单,改单"'
    for address in ("A4", "C4", "B3", "B7", "B9", "B10", "B12", "B13"):
        assert _rules(sheet, address) == [original]
    for address in ("B4", "B6", "B8", "B11"):
        rule, = _rules(sheet, address)
        assert rule is not original and "备料单" in rule.formula1
    for name in SHEETS[1:]:
        assert not _rules(workbook[name], "B7")
        assert not _rules(workbook[name], "B9")
        assert not _rules(workbook[name], "B10")


def test_literal_option_extension_is_idempotent_after_save_and_reload():
    workbook = _workbook()
    for sheet in workbook:
        _add(sheet)
    _apply(workbook)
    stream = BytesIO()
    workbook.save(stream)
    restored = openpyxl.load_workbook(BytesIO(stream.getvalue()))
    before = _validation_snapshot(restored)
    _apply(restored)
    assert _validation_snapshot(restored) == before


@pytest.mark.parametrize("formula", ["=OrderTypes", "'原 选项'!$A$1:$A$3"])
def test_named_or_fixed_range_uses_live_separate_list_and_leaves_source_consumers_untouched(formula):
    workbook = _workbook()
    source = workbook.create_sheet("原 选项")
    for row, value in enumerate(("正式PO", "备料", '=IF(1=1,"加单","")'), 1):
        source.cell(row, 1, value)
    workbook.defined_names.add(DefinedName("OrderTypes", attr_text="'原 选项'!$A$1:$A$3"))
    original_values = _snapshot(workbook)
    for name in SHEETS:
        _add(workbook[name], formula)
    other = _add(source, "=OrderTypes", "D1:D8")
    _apply(workbook)
    assert workbook.defined_names["OrderTypes"].attr_text == "'原 选项'!$A$1:$A$3"
    assert other.formula1 == "=OrderTypes" and str(other.sqref) == "D1:D8"
    assert {name: _snapshot(workbook)[name] for name in original_values} == original_values
    helpers = [s for s in workbook if s.title not in original_values]
    assert len(helpers) == 1 and helpers[0].sheet_state == "veryHidden"
    helper = helpers[0]
    assert helper["B2"].value == '=IF(\'原 选项\'!$A$1="","",\'原 选项\'!$A$1)'
    assert helper["B4"].value == '=IF(\'原 选项\'!$A$3="","",\'原 选项\'!$A$3)'
    assert {c.value for c in helper["B"] if c.data_type != "f"} >= set(ORDER_TYPES)
    for name in SHEETS:
        assert _rules(workbook[name], "B4")[0].formula1.startswith("=_HX360_OrderTypes_")
    before = (_snapshot(workbook), _validation_snapshot(workbook), list(workbook.defined_names))
    _apply(workbook)
    assert (_snapshot(workbook), _validation_snapshot(workbook), list(workbook.defined_names)) == before


def test_local_named_range_takes_precedence_and_full_existing_list_keeps_formula():
    workbook = _workbook()
    for sheet in workbook:
        for row, value in enumerate(ORDER_TYPES, 1):
            sheet.cell(row, 30, value)
        sheet.defined_names.add(DefinedName("LocalTypes", attr_text=f"'{sheet.title}'!$AD$1:$AD$5"))
        _add(sheet, "=LocalTypes")
    _apply(workbook)
    assert workbook.sheetnames == list(SHEETS)
    for sheet in workbook:
        assert _rules(sheet, "B4")[0].formula1 == "=LocalTypes"


def test_named_original_four_options_do_not_grow_after_repeated_save_reload():
    workbook = _workbook()
    source = workbook.create_sheet("选项")
    for row, value in enumerate(ORDER_TYPES[:-1], 1):
        source.cell(row, 1, value)
    workbook.defined_names.add(DefinedName("OrderTypes", attr_text="'选项'!$A$1:$A$4"))
    for name in SHEETS:
        _add(workbook[name], "=OrderTypes")
    _apply(workbook)
    before = (_snapshot(workbook), _validation_snapshot(workbook), list(workbook.defined_names))
    for _ in range(3):
        stream = BytesIO()
        workbook.save(stream)
        workbook = openpyxl.load_workbook(BytesIO(stream.getvalue()))
        _apply(workbook)
        assert (_snapshot(workbook), _validation_snapshot(workbook), list(workbook.defined_names)) == before


def _entry_source():
    from app.services import customer_order_unified as service
    workbook = _workbook()
    for name in SHEETS:
        sheet = workbook[name]
        headers = service.ITEM_HEADERS if name == "ITEM表" else service.ORDER_HEADERS
        for column, value in enumerate(headers, 1):
            sheet.cell(3, column, value)
        reference, product, inspection = (5, 8, 25) if name == "ITEM表" else (4, 7, 18)
        for row, when in ((4, date(2020, 1, 1)), (5, date(2099, 1, 1))):
            sheet.cell(row, reference, f"RL-100-{row}")
            sheet.cell(row, product, f"OLD-{row}")
            sheet.cell(row, inspection, when)
            sheet.cell(row, 2, "备料单" if row == 4 else "历史自定义类型")
        sheet["A16"] = "已走货订单"
        sheet["B17"] = "已走货人工类型"
        _add(sheet)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _assert_entry_result(workbook):
    for name in SHEETS:
        sheet = workbook[name]
        reference = 5 if name == "ITEM表" else 4
        historical = {c.value: c.row for c in sheet._cells.values()
                      if c.column == reference and str(c.value).startswith("RL-100-")}
        assert sheet.cell(historical["RL-100-4"], 2).value == "备料单"
        assert sheet.cell(historical["RL-100-5"], 2).value == "历史自定义类型"
        for row in historical.values():
            rule, = _rules(sheet, f"B{row}")
            assert set(ORDER_TYPES) <= set(rule.formula1.strip('"').split(","))
        completed = next(c.row for c in sheet._cells.values() if c.value == "已完成（按验货日期）")
        assert completed > 17 and historical["RL-100-4"] > completed
        assert not _rules(sheet, f"B{completed}")
        assert sheet["B13"].value == "人工取消类型"
        assert sheet["B17"].value == "已走货人工类型"
        assert not _rules(sheet, "B13") and not _rules(sheet, "B17")


def test_standalone_organize_entry_covers_moved_completed_rows_and_is_repeatable():
    from app.services.customer_order_360_workbook import organize_workbook
    output, _, report, _ = organize_workbook("source.xlsx", _entry_source())
    workbook = openpyxl.load_workbook(BytesIO(output))
    _assert_entry_result(workbook)
    assert report["order_type_dropdowns"]["applied"]
    before = (_snapshot(workbook), _validation_snapshot(workbook))
    repeated, _, _, _ = organize_workbook("organized.xlsx", output)
    workbook = openpyxl.load_workbook(BytesIO(repeated))
    _assert_entry_result(workbook)
    assert (_snapshot(workbook), _validation_snapshot(workbook)) == before


def test_normal_export_entry_covers_new_and_moved_orders(monkeypatch):
    from app.services import customer_order_unified as service
    from app.services.huakang_a_order_legacy import schedule_parser
    record = {
        "contract_no": "RL-900-1", "customer_release": "OH-9001", "customer_po": "PO-900",
        "item_no": "NEW", "description": "New toy", "quantity": 120, "master_carton_qty": 6,
        "planned_inspection_date": "2099-01-01", "factory_commit_date": "2099-01-08",
        "parse_ok": True, "parse_warnings": [],
    }
    monkeypatch.setattr(schedule_parser, "parse_po_file", lambda *_: dict(record))
    output, _, _ = service.export_unified_customer_schedule(
        customer_code="360", factory_id="huakang-a", received_date="2026-10-19",
        po_files=[("new.pdf", b"fixture")], schedule_file_name="source.xlsx", schedule_content=_entry_source(),
    )
    workbook = openpyxl.load_workbook(BytesIO(output))
    _assert_entry_result(workbook)
    for name in SHEETS:
        sheet = workbook[name]
        reference = 5 if name == "ITEM表" else 4
        row = next(c.row for c in sheet._cells.values() if c.column == reference and c.value == "RL-900-1")
        assert "备料单" in _rules(sheet, f"B{row}")[0].formula1


@pytest.mark.parametrize("formula", [
    "=INDIRECT(A1)", "=OFFSET(ITEM表!$AD$1,0,0,4,1)",
    "'[external.xlsx]Sheet1'!$A$1:$A$4", "ITEM表!$AD$1:$AE$3",
])
def test_unsupported_lists_fail_before_any_sheet_is_changed(formula):
    workbook = _workbook()
    _add(workbook["ITEM表"])
    _add(workbook["接单表"], formula)
    before = (_snapshot(workbook), _validation_snapshot(workbook))
    with pytest.raises(OrderTypeValidationError, match="固定命名区域"):
        _apply(workbook)
    assert (_snapshot(workbook), _validation_snapshot(workbook)) == before


def test_unrelated_factory_and_customer_are_noops():
    workbook = _workbook()
    for sheet in workbook:
        _add(sheet)
    before = deepcopy(_validation_snapshot(workbook))
    for factory, customer in (("huaxing", "360"), ("huakang-a", "green-toys")):
        result = ensure_360_order_type_dropdowns(workbook, factory_id=factory, customer_code=customer)
        assert not result["applied"]
    assert _validation_snapshot(workbook) == before


def test_overlap_is_reported_without_duplicate_dropdowns():
    workbook = _workbook()
    _add(workbook["ITEM表"], cells="B4:B8")
    _add(workbook["ITEM表"], cells="B7:B11")
    before = _validation_snapshot(workbook)
    with pytest.raises(OrderTypeValidationError, match="重叠验证"):
        _apply(workbook)
    assert _validation_snapshot(workbook) == before


def test_real_360_source_values_formulas_and_original_remain_unchanged():
    manifest = require_private_manifest()
    source = manifest["schedule"]
    original = source.read_bytes()
    workbook = openpyxl.load_workbook(BytesIO(original))
    before = _snapshot(workbook)
    _apply(workbook)
    assert _snapshot(workbook) == before
    checks = manifest["validation_checks"]
    assert {entry["sheet"] for entry in checks} == set(SHEETS)
    for entry in checks:
        sheet = workbook[entry["sheet"]]
        assert entry["include"] and entry["exclude"]
        for coordinate in entry["include"]:
            rule, = _rules(sheet, coordinate)
            assert set(ORDER_TYPES) <= set(rule.formula1.strip('"').split(","))
        for coordinate in entry["exclude"]:
            assert not _rules(sheet, coordinate)
        for coordinate, expected in entry["preserve_values"].items():
            assert sheet[coordinate].value == expected
    assert sha256(source.read_bytes()).digest() == sha256(original).digest()
