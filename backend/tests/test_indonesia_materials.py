"""Business-flow regressions using isolated synthetic records only."""
import os
os.environ.setdefault('DATABASE_URL', 'sqlite://')

import base64
from copy import deepcopy
from dataclasses import replace
from decimal import Decimal
from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile

import openpyxl
import pytest
from PIL import Image
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from lxml import etree as ET
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from app.api import indonesia_materials as api
from app.db import get_db
from app.models.indonesia_materials import IndonesiaMaterialRecord as Record, IndonesiaMaterialCommand as Receipt, IndonesiaMaterialAsset as Asset
from app.models.auth import AuthIamState
from app.services import indonesia_materials as service
from app.services.auth import AuthContext, AuthGrantContext, AuthOverrideContext, get_current_user
from app.services.indonesia_material_calculations import packaging_totals, production_plan
from app.services.indonesia_material_excel import TEMPLATE, NS, DRAWING_PART, SHEET_PART, SHEET_NAME, CALC_NAME
from test_indonesia_material_export import detail_book
from app.services.permission_codes import INDONESIA_MATERIAL_PERMISSION_CODES

BASE = '/api/indonesia-materials'


def actor(*actions, factory='huaxing', department='sales-business', role='synthetic'):
    codes = frozenset('indonesia_material:' + a for a in actions)
    grant = AuthGrantContext(role_id=role, role_name='测试角色', factory_id=factory, department=department, permissions=codes)
    return AuthContext(id='synthetic-user', username='synthetic', display_name='合成测试', roles=('test',), role_codes=('test',),
                       permissions=codes, factory_scopes=(factory,), department_scopes=(department,), grants=(grant,),
                       active_permission_codes=frozenset(INDONESIA_MATERIAL_PERMISSION_CODES))


@pytest.fixture
def client(tmp_path, monkeypatch):
    engine = create_engine('sqlite:///' + (tmp_path / 'materials.db').as_posix(),
                           connect_args={'check_same_thread': False, 'timeout': 20})
    Record.metadata.create_all(engine, tables=list(service.TABLES))
    AuthIamState.__table__.create(engine)
    with Session(engine) as db:
        db.add(AuthIamState(key='identity_mutation_lock', value_json='{}'))
        db.commit()
    auth = {'user': actor('read', 'write', 'receipt_confirm', 'export', role='position_sales_supervisor')}
    def session():
        with Session(engine, expire_on_commit=False) as db:
            yield db
    # Parse behavior is separately covered; API tests do not launch subprocesses.
    async def run_job(function, *args, db=None, **kwargs):
        if db:
            db.rollback()
        return function(*args, **kwargs)
    monkeypatch.setattr(api, 'run_order_job', run_job)
    monkeypatch.setattr(api, 'refresh_user', lambda request, db: auth['user'])
    app = FastAPI(); app.include_router(api.router)
    app.dependency_overrides[get_db] = session
    app.dependency_overrides[get_current_user] = lambda: auth['user']
    with TestClient(app) as test:
        test.engine, test.materials_auth = engine, auth
        yield test
    engine.dispose()


def command(client, action, data, row=None, *, factory='huaxing', status=200, operation_id=None):
    if action == 'option_save' and row:
        data = {key: value for key, value in data.items() if key in service.Option.model_fields}
    shared = action.startswith('option_')
    body = dict(factory_id='shared-options' if shared else factory, operation_id=operation_id or uuid4().hex, action=action, data=data,
                entity_id=row['id'] if row else None, expected_version=row['version'] if row else 0)
    response = client.post(BASE + ('/shared-options/commands' if shared else '/commands'), json=body)
    assert response.status_code == status, response.text
    return response.json()['record'] if status == 200 else response


def material(client, code='00001', **changes):
    image = BytesIO(); Image.new('RGB', (4, 4), 'green').save(image, 'PNG')
    data = dict(code=code, name='合成螺丝', english_name='Synthetic screw', specification='TEST', category='辅助材料',
                unit='PCS', customs_unit='KGM', net_weight_kg='.001', billing_weight_kg='.002',
                price='.10', billing_unit='PCS', currency='USD', procurement_days=18,
                supplier='合成供应商', usage='合成测试装配', department='合成装配部门', export_company='合成出口公司',
                image_data='data:image/png;base64,' + base64.b64encode(image.getvalue()).decode())
    return command(client, 'material_save', data | changes)


def customer(client, name='合成客户', **changes):
    return command(client, 'customer_save', dict(name=name, status='active', notes='') | changes)


def ensure_customer(client, name='合成客户'):
    existing = next((r for r in workspace(client)['records'] if r['kind'] == 'customer' and r['data']['name'] == name), None)
    return existing or customer(client, name)


def inline_material(code, **changes):
    return dict(code=code, name='货号内新增物料', unit='PCS', procurement_days=18,
                net_weight_kg='0.00123456', supplier='货号内供应商', price='0.12') | changes


def test_shared_options_ordinary_sales_create_but_only_active_supervisors_manage(client):
    rule = command(client, 'option_save', dict(group='customs_unit', value='CUSTOM', billing_basis='weight', kg_per_unit='10'))
    client.materials_auth['user'] = actor('read', 'write')
    plain = command(client, 'option_save', dict(group='category', value='普通业务共同新增'))
    assert plain['scope'] == 'shared-options'
    command(client, 'option_save', plain['data'] | {'status': 'inactive'}, plain, status=403)
    command(client, 'option_delete', {}, plain, status=403)
    command(client, 'option_save', rule['data'] | {'kg_per_unit': '100'}, rule, status=403)
    command(client, 'option_save', dict(group='customs_unit', value='NEW-WEIGHT', billing_basis='weight', kg_per_unit='5'), status=403)
    user = actor('read', 'write', role='position_sales_supervisor')
    client.materials_auth['user'] = replace(user, grants=(replace(user.grants[0], valid_until='2020-01-01T00:00:00Z'),
                                                           actor('read', 'write').grants[0]))
    command(client, 'option_delete', {}, plain, status=403)
    client.materials_auth['user'] = user
    inactive = command(client, 'option_save', plain['data'] | {'status': 'inactive'}, plain)
    assert inactive['data']['status'] == 'inactive'


def test_shared_write_requires_sales_read_despite_pmc_read_and_explicit_write(client):
    user = actor('read', 'write')
    deny = AuthOverrideContext(id='deny-sales-read', permission_code='indonesia_material:read', effect='deny', factory_id='huaxing', department='sales-business')
    warehouse = actor('read', department='pmc-warehouse').grants[0]
    client.materials_auth['user'] = replace(user, grants=(*user.grants, warehouse), overrides=(deny,))
    assert client.get(BASE + '/shared-options').status_code == 200
    command(client, 'option_save', dict(group='supplier', value='不能借仓库查看权限维护'), status=403)


@pytest.mark.parametrize('group', ['category', 'customs_unit', 'unit', 'usage', 'currency', 'supplier', 'export_company'])
def test_deleted_global_choice_cannot_be_referenced_by_a_stale_new_material_form(client, group):
    choice = command(client, 'option_save', dict(group=group, value='STALE'))
    command(client, 'option_delete', {}, choice)
    client.materials_auth['user'] = actor('read', 'write', factory='huadeng')
    command(client, 'material_save', inline_material('STALE-MATERIAL', **{group: 'STALE'}), factory='huadeng', status=409)
    assert all(row['kind'] == 'option' for row in client.get(BASE + '/workspace?factory_id=huadeng').json()['records'])


@pytest.mark.parametrize('group', ['currency', 'supplier'])
def test_shared_inactive_choice_is_rechecked_for_new_purchase_but_old_input_stays_frozen(client, group):
    choice = command(client, 'option_save', dict(group=group, value='FROZEN'))
    m = material(client)
    old = supply(client, m, **{group: 'FROZEN'})
    command(client, 'option_delete', {}, choice, status=409)
    command(client, 'option_save', choice['data'] | {'status': 'inactive'}, choice)
    supply_payload = {key: value for key, value in old['data'].items() if key in service.s.Supply.model_fields}
    command(client, 'supply_save', supply_payload | {'batch': 'NEW'}, status=409)
    edited = command(client, 'supply_save', supply_payload | {'normal_quantity': '89'}, old)
    assert edited['data']['material_snapshot'] == old['data']['material_snapshot']


def test_used_shared_choice_checks_nested_history_without_exposing_other_factory_data(client):
    choice = command(client, 'option_save', dict(group='category', value='旧类别'))
    m = material(client, category='旧类别'); product(client, m); order(client)
    command(client, 'material_save', m['data'] | {'category': '后来的类别'}, m)
    client.materials_auth['user'] = actor('read', 'write', factory='huadeng', role='position_sales_supervisor')
    response = command(client, 'option_delete', {}, choice, factory='huadeng', status=409)
    assert 'TEST-A' not in response.text and '合成客户' not in response.text
    assert client.get(BASE + '/workspace?factory_id=huaxing').status_code == 403


def test_legacy_factory_option_receipt_replays_but_uncommitted_requests_never_become_global(client):
    from app.models.indonesia_materials import IndonesiaMaterialOption as PrivateOption
    body = dict(factory_id='huaxing', operation_id=uuid4().hex, action='option_save', entity_id=None, expected_version=0,
                data=dict(group='currency', value='OLD'))
    data = dict(group='currency', value='OLD', status='active', billing_basis='quantity', kg_per_unit='1')
    with Session(client.engine) as db:
        row = PrivateOption(id='legacy-option', factory_id='huaxing', business_key='legacy-key', version=1, data=data, updated_at='old')
        db.add(row); db.flush()
        result = dict(record=service.view(row))
        db.add(Receipt(operation_id=body['operation_id'], factory_id='huaxing', actor_id='synthetic-user', action='option_save',
                       fingerprint=service.digest(service.s.Command.model_validate(body).model_dump(mode='json')), changes=[], result=result, created_at='old'))
        db.commit()
    assert client.post(BASE + '/commands', json=body).json() == result
    body['operation_id'] = uuid4().hex
    assert client.post(BASE + '/commands', json=body).status_code == 409
    assert not client.get(BASE + '/shared-options').json()['records']


def test_options_are_shared_versioned_and_private_records_stay_factory_scoped(client):
    defaults = workspace(client)['option_defaults']
    assert {'KGM', 'TNE', 'PCE'} <= {item['value'] for item in defaults if item['group'] == 'customs_unit'}
    assert len([item for item in defaults if item['group'] == 'supplier']) == 8
    row = command(client, 'option_save', dict(group='currency', value='hkd'))
    assert row['data']['value'] == 'HKD'
    row = command(client, 'option_save', row['data'] | {'status': 'inactive'}, row)
    assert row['version'] == 2
    command(client, 'option_save', row['data'] | {'value': 'EUR'}, row, status=422)
    command(client, 'option_save', dict(group='currency', value='HKD'), status=409)
    client.materials_auth['user'] = actor('read', 'write', factory='huadeng', role='position_sales_supervisor')
    updated = command(client, 'option_save', row['data'] | {'status': 'active'}, row, factory='huadeng')
    assert client.get(BASE + '/workspace', params={'factory_id': 'huadeng'}).json()['records'] == [updated]
    client.materials_auth['user'] = actor('read')
    command(client, 'option_save', dict(group='supplier', value='禁止越权'), status=403)


@pytest.mark.parametrize('group,value', [('category', '塑胶'), ('customs_unit', 'KGM'), ('unit', 'pcs'),
                                      ('usage', '产品配件'), ('currency', 'USD'),
                                      ('supplier', '华兴玩具制品（河源）有限公司'), ('export_company', '新利達環球發展有限公司')])
def test_default_options_can_be_deleted_and_restored_idempotently(client, group, value):
    identity = dict(group=group, value=value)
    op = uuid4().hex
    removed = command(client, 'option_delete', identity, operation_id=op)
    assert command(client, 'option_delete', identity, operation_id=op) == removed
    assert removed['version'] == 1 and removed['data']['deleted']
    assert removed['data']['deleted_by'] == 'synthetic-user'
    command(client, 'option_delete', identity, status=409)
    command(client, 'option_save', identity, status=409)
    command(client, 'option_save', identity, removed, status=404)
    assert next(row for row in workspace(client)['records'] if row['id'] == removed['id']) == removed
    op = uuid4().hex
    restored = command(client, 'option_restore', {}, removed, operation_id=op)
    assert command(client, 'option_restore', {}, removed, operation_id=op) == restored
    assert restored['id'] == removed['id'] and not restored['data']['deleted']
    assert all(restored['data'][key] == removed['data'][key] for key in ('group', 'value', 'status', 'billing_basis', 'kg_per_unit'))
    command(client, 'option_delete', {}, removed, status=409)  # stale version cannot re-delete
    editable = {key: restored['data'][key] for key in service.Option.model_fields}
    saved = command(client, 'option_save', editable | {'status': 'inactive'}, restored)
    assert saved['version'] == 3 and saved['data']['status'] == 'inactive'


def test_custom_option_archive_validates_identity_payload_permissions_and_factory(client):
    row = command(client, 'option_save', dict(group='currency', value='RMB'))
    command(client, 'option_delete', {}, status=422)
    command(client, 'option_delete', dict(group='currency', value='UNKNOWN'), status=422)
    command(client, 'option_delete', {'value': 'USD'}, row, status=422)
    command(client, 'option_delete', dict(group='currency', value='USD', status='inactive'), status=422)
    default_override = command(client, 'option_save', dict(group='currency', value='USD', status='inactive'))
    command(client, 'option_delete', dict(group='currency', value='USD'), status=409)
    assert next(r for r in workspace(client)['records'] if r['id'] == default_override['id']) == default_override
    deleted = command(client, 'option_delete', {}, row)
    command(client, 'option_restore', {'value': 'EUR'}, deleted, status=422)
    client.materials_auth['user'] = actor('read', 'write', factory='huadeng')
    command(client, 'option_restore', {}, deleted, factory='huadeng', status=403)
    command(client, 'option_delete', {}, default_override, factory='huadeng', status=403)
    assert len(client.get(BASE + '/workspace', params={'factory_id': 'huadeng'}).json()['records']) == 2
    client.materials_auth['user'] = actor('read')
    command(client, 'option_delete', {}, default_override, status=403)
    command(client, 'option_delete', dict(group='currency', value='HKD'), status=403)
    command(client, 'option_restore', {}, deleted, status=403)


def test_deleted_customs_default_does_not_reappear_or_reprice_existing_purchase(client):
    m = material(client, billing_mode='automatic')
    product(client, m); o = order(client)
    sup = supply(client, m, normal='100', spare='0', billing_mode='automatic')
    command(client, 'option_delete', dict(group='customs_unit', value='KGM'), status=409)
    removed = command(client, 'option_save', dict(group='customs_unit', value='KGM', status='inactive', billing_basis='weight'))
    m = command(client, 'material_save', m['data'] | {'english_name': 'Historical description corrected'}, m)
    assert m['data']['customs_rule'] == dict(basis='weight', divisor='1')
    command(client, 'material_save', inline_material('NEW', billing_mode='automatic'), status=409)
    # New purchases must see the tombstone instead of falling back to built-in KGM.
    command(client, 'supply_save', {key: value for key, value in sup['data'].items() if key in service.s.Supply.model_fields}, status=422)
    assert next(r for r in workspace(client)['records'] if r['id'] == sup['id']) == sup
    sup = command(client, 'receipt_confirm', {'actual_date': '2026-10-22'}, sup)
    alloc = allocation(client, sup, o, normal='100', spare='0')
    sh = command(client, 'shipment_confirm', {}, shipment(client, alloc, normal='100', spare='0'))
    exported = client.get(BASE + f"/shipments/{sh['id']}/export", params={'factory_id': 'huaxing'}).content
    restored = command(client, 'option_save', removed['data'] | {'status': 'active'}, removed)
    assert restored['data']['status'] == 'active'
    assert client.get(BASE + f"/shipments/{sh['id']}/export", params={'factory_id': 'huaxing'}).content == exported
    material(client, code='NEW', billing_mode='automatic')


@pytest.mark.parametrize('unit,basis,factor,amount', [('KGM','weight','1','2000'), ('TNE','weight','1000','2'), ('PCE','quantity','1','1000'), ('METRICTON','weight','1000','2')])
def test_automatic_customs_unit_pricing_reaches_frozen_excel(client, unit, basis, factor, amount):
    if unit == 'METRICTON':
        command(client, 'option_save', dict(group='customs_unit', value=unit, billing_basis='weight', kg_per_unit='1000'))
    m = material(client, unit='pcs', customs_unit=unit, billing_mode='automatic', net_weight_kg='2', billing_weight_kg='999',
                 billing_basis='quantity', billing_factor='99', billing_unit='WRONG')
    assert (m['data']['billing_basis'], m['data']['billing_factor'], m['data']['billing_unit']) == (basis, factor, unit)
    assert m['data']['billing_weight_kg'] == m['data']['net_weight_kg'] == '2'
    product(client, m); o = order(client)
    sup = supply(client, m, normal='100', spare='0', billing_mode='automatic', price='10', billing_factor='77')
    assert (sup['data']['billing_basis'], sup['data']['billing_factor']) == (basis, factor)
    if unit == 'METRICTON':
        rule = next(row for row in workspace(client)['records'] if row['kind'] == 'option')
        command(client, 'option_save', rule['data'] | {'kg_per_unit': '2000'}, rule)
    sup = command(client, 'receipt_confirm', {'actual_date': '2026-10-22'}, sup)
    alloc = allocation(client, sup, o, normal='100', spare='0')
    sh = shipment(client, alloc, normal='100', spare='0')
    sh = command(client, 'shipment_confirm', {}, sh)
    response = client.get(BASE + f"/shipments/{sh['id']}/export", params={'factory_id': 'huaxing'})
    sheet = detail_book(BytesIO(response.content), data_only=True).worksheets[0]
    assert Decimal(str(sheet['AP4'].value)) == Decimal(amount)
    assert Decimal(str(sheet['K4'].value)) == (Decimal('200') / Decimal(factor) if basis == 'weight' else Decimal('100'))


def test_auto_pricing_uses_maintained_unit_rules_without_rewriting_old_purchase(client):
    rule = command(client, 'option_save', dict(group='customs_unit', value='METRICTON', billing_basis='weight', kg_per_unit='1000'))
    m = material(client, customs_unit='METRICTON', billing_mode='automatic')
    old = supply(client, m, billing_mode='automatic')
    command(client, 'option_save', rule['data'] | {'kg_per_unit': '2000'}, rule)
    assert next(row for row in workspace(client)['records'] if row['id'] == old['id']) == old
    data = {key: value for key, value in old['data'].items() if key in service.s.Supply.model_fields}
    changed = command(client, 'supply_save', data | {'billing_factor': '999', 'billing_basis': 'quantity'}, old)
    assert changed['data']['billing_factor'] == '1000' and changed['data']['billing_basis'] == 'weight'


def test_inactive_or_unknown_customs_unit_blocks_automatic_pricing(client):
    command(client, 'option_save', dict(group='customs_unit', value='KGM', status='inactive', billing_basis='weight'))
    command(client, 'material_save', inline_material('BLOCKED', customs_unit='KGM', billing_mode='automatic'), status=409)
    command(client, 'material_save', inline_material('UNKNOWN', customs_unit='UNKNOWN', billing_mode='automatic'), status=422)
    # Legacy API data is preserved without opting it into a new calculation.
    legacy = material(client, customs_unit='UNKNOWN', billing_factor='7', billing_unit='HISTORICAL')
    assert legacy['data']['billing_factor'] == '7' and legacy['data']['billing_unit'] == 'HISTORICAL'


def test_customer_and_product_delete_restore_preserve_materials_identity_and_idempotency(client):
    m = material(client); p = product(client, m)
    c = next(row for row in workspace(client)['records'] if row['kind'] == 'customer')
    command(client, 'customer_delete', {}, c, status=409)
    op = uuid4().hex
    removed = command(client, 'product_delete', {}, p, operation_id=op)
    assert command(client, 'product_delete', {}, p, operation_id=op) == removed
    assert removed['data']['deleted'] and removed['data']['materials'] == p['data']['materials']
    deleted_customer = command(client, 'customer_delete', {}, c)
    command(client, 'product_restore', {}, removed, status=422)
    command(client, 'customer_restore', {}, deleted_customer)
    restored = command(client, 'product_restore', {}, removed)
    assert not restored['data']['deleted'] and restored['id'] == p['id']
    assert next(row for row in workspace(client)['records'] if row['id'] == m['id']) == m


def test_referenced_product_cannot_be_deleted_or_edited_via_archive_payload(client):
    m = material(client); p = product(client, m); order(client)
    before = workspace(client)['records']
    command(client, 'product_delete', {}, p, status=409)
    command(client, 'product_delete', {'materials': []}, p, status=422)
    assert workspace(client)['records'] == before
    client.materials_auth['user'] = actor('read', 'write', factory='huadeng')
    command(client, 'product_delete', {}, p, factory='huadeng', status=404)
    client.materials_auth['user'] = actor('read')
    command(client, 'product_delete', {}, p, status=403)


def test_material_deletion_archives_only_material_and_unlinks_parent_atomically(client):
    m = material(client); keep = material(client, 'KEEP')
    p = product(client, m, materials=[dict(material_id=m['id'], usage='2'), dict(material_id=keep['id'], usage='3')])
    payload = p['data'] | dict(materials=[dict(material_id=keep['id'], usage='3')],
                              material_deletions=[dict(material_id=m['id'], expected_version=m['version'])])
    op = uuid4().hex
    saved = command(client, 'product_save', payload, p, operation_id=op)
    assert command(client, 'product_save', payload, p, operation_id=op) == saved
    assert saved['id'] == p['id'] and saved['version'] == p['version'] + 1
    assert saved['data']['code'] == p['data']['code'] and not saved['data'].get('deleted')
    assert saved['data']['materials'] == [dict(material_id=keep['id'], usage='3')]
    assert 'material_deletions' not in saved['data']
    rows = workspace(client)['records']
    removed = next(r for r in rows if r['id'] == m['id'])
    assert removed['data']['deleted'] and removed['data']['deleted_by'] == 'synthetic-user'
    assert next(r for r in rows if r['id'] == keep['id']) == keep
    with Session(client.engine) as db:
        assert [c['after']['kind'] for c in db.get(Receipt, op).changes] == ['material', 'product']
    command(client, 'material_save', m['data'], removed, status=404)
    command(client, 'material_save', m['data'], status=409)  # archived code stays reserved
    command(client, 'product_save', saved['data'] | {'materials': [dict(material_id=m['id'], usage='1')]}, saved, status=404)
    command(client, 'material_restore', {'name': '不得覆盖'}, removed, status=422)
    restore_op = uuid4().hex
    restored = command(client, 'material_restore', {}, removed, operation_id=restore_op)
    assert command(client, 'material_restore', {}, removed, operation_id=restore_op) == restored
    assert restored['id'] == m['id'] and not restored['data']['deleted']
    assert {k: restored['data'][k] for k in m['data']} == m['data']
    assert next(r for r in workspace(client)['records'] if r['id'] == saved['id']) == saved
    command(client, 'material_restore', {}, removed, status=409)


@pytest.mark.parametrize('archived', [False, True])
def test_deleting_shared_material_refuses_other_products_including_archived(client, archived):
    m = material(client); p = product(client, m); other = product(client, m, code='OTHER')
    if archived:
        command(client, 'product_delete', {}, other)
    before = workspace(client)['records']
    command(client, 'product_save', p['data'] | dict(materials=[], material_deletions=[
        dict(material_id=m['id'], expected_version=m['version'])]), p, status=409)
    assert workspace(client)['records'] == before
    # Removing this product's relation alone neither archives shared material nor changes the other product.
    command(client, 'product_save', p['data'] | {'materials': []}, p)
    assert next(r for r in workspace(client)['records'] if r['id'] == m['id']) == m


@pytest.mark.parametrize('source', ['order', 'supply'])
def test_material_with_business_history_cannot_be_deleted(client, source):
    m = material(client); p = product(client, m)
    history = order(client) if source == 'order' else supply(client, m)
    p = command(client, 'product_save', p['data'] | {'materials': []}, p)
    before = workspace(client)['records']
    command(client, 'product_save', p['data'] | {'material_deletions': [
        dict(material_id=m['id'], expected_version=m['version'])]}, p, status=409)
    assert workspace(client)['records'] == before
    assert next(r for r in before if r['id'] == history['id']) == history


def test_material_deletion_conflict_rolls_back_new_material_parent_and_earlier_deletion(client):
    first = material(client, 'FIRST'); stale = material(client, 'STALE'); p = product(client, first)
    command(client, 'material_save', stale['data'] | {'name': '别人已更新'}, stale)
    before = workspace(client)['records']
    command(client, 'product_save', p['data'] | dict(name='不能部分保存', materials=[
        dict(material=inline_material('ROLLBACK-NEW'), usage='1')], material_deletions=[
            dict(material_id=first['id'], expected_version=first['version']),
            dict(material_id=stale['id'], expected_version=stale['version'])]), p, status=409)
    assert workspace(client)['records'] == before


@pytest.mark.parametrize('mode', ['duplicate', 'retained', 'zero_version', 'extra_field', 'missing'])
def test_material_deletion_requires_distinct_original_ids_and_versions(client, mode):
    m = material(client); p = product(client, m)
    deletion = dict(material_id=m['id'], expected_version=m['version'])
    if mode == 'zero_version':
        deletion['expected_version'] = 0
    if mode == 'extra_field':
        deletion['name'] = '不能覆盖'
    if mode == 'missing':
        deletion['material_id'] = 'missing'
    payload = p['data'] | dict(materials=p['data']['materials'] if mode == 'retained' else [],
                              material_deletions=[deletion, deletion] if mode == 'duplicate' else [deletion])
    before = workspace(client)['records']
    command(client, 'product_save', payload, p, status=404 if mode == 'missing' else 422)
    assert workspace(client)['records'] == before


def test_material_deletion_and_restore_enforce_factory_permissions_and_parent_version(client):
    m = material(client); p = product(client, m)
    payload = p['data'] | dict(materials=[], material_deletions=[dict(material_id=m['id'], expected_version=m['version'])])
    newer = command(client, 'product_save', p['data'] | {'name': '新版货号'}, p)
    command(client, 'product_save', payload, p, status=409)
    client.materials_auth['user'] = actor('read')
    command(client, 'product_save', payload, newer, status=403)
    client.materials_auth['user'] = actor('read', 'write', factory='huadeng')
    command(client, 'customer_save', dict(name='合成客户'), factory='huadeng')
    other = command(client, 'product_save', dict(code='OTHER', customer='合成客户', materials=[]), factory='huadeng')
    command(client, 'product_save', other['data'] | {'material_deletions': payload['material_deletions']}, other, factory='huadeng', status=404)
    client.materials_auth['user'] = actor('read', 'write')
    command(client, 'product_save', payload, newer)
    removed = next(r for r in workspace(client)['records'] if r['id'] == m['id'])
    client.materials_auth['user'] = actor('read', 'write', factory='huadeng')
    command(client, 'material_restore', {}, removed, factory='huadeng', status=404)
    client.materials_auth['user'] = actor('read')
    command(client, 'material_restore', {}, removed, status=403)


def test_product_creates_materials_and_recipe_in_one_idempotent_command(client):
    customer(client)
    data = dict(code='INLINE', customer='合成客户', daily_capacity='5000', materials=[
        dict(material=inline_material('INLINE-A'), usage='2.12345678'),
        dict(material=inline_material('INLINE-B'), usage='3')])
    operation = uuid4().hex
    p = command(client, 'product_save', data, operation_id=operation)
    assert command(client, 'product_save', data, operation_id=operation) == p
    rows = workspace(client)['records']
    materials = [r for r in rows if r['kind'] == 'material']
    assert len(materials) == 2
    assert {m['data']['code'] for m in materials} == {'INLINE-A', 'INLINE-B'}
    assert p['data']['materials'] == [dict(material_id=materials_by_code(materials, 'INLINE-A')['id'], usage='2.12345678'),
                                     dict(material_id=materials_by_code(materials, 'INLINE-B')['id'], usage='3')]
    with Session(client.engine) as db:
        receipt = db.get(Receipt, operation)
        assert [c['after']['kind'] for c in receipt.changes] == ['material', 'material', 'product']


def materials_by_code(rows, code):
    return next(r for r in rows if r['kind'] == 'material' and r['data']['code'] == code)


def test_product_rolls_back_earlier_materials_on_duplicate_code(client):
    customer(client)
    existing = material(client, 'TAKEN')
    before = workspace(client)['records']
    command(client, 'product_save', dict(code='INLINE', customer='合成客户', materials=[
        dict(material=inline_material('WOULD-BE-ORPHAN'), usage='1'),
        dict(material=inline_material(existing['data']['code']), usage='2')]), status=409)
    assert workspace(client)['records'] == before


def test_product_stale_material_version_rolls_back_new_material_and_parent(client):
    m = material(client); p = product(client, m)
    command(client, 'material_save', m['data'] | {'supplier': '另一业务已修改'}, m)
    before = workspace(client)['records']
    command(client, 'product_save', p['data'] | {'materials': [
        dict(material=inline_material('NEW-BEFORE-CONFLICT'), usage='1'),
        dict(material_id=m['id'], material_version=m['version'], material=m['data'] | {'supplier': '旧值覆盖'}, usage='4')
    ]}, p, status=409)
    assert workspace(client)['records'] == before


def test_product_nested_edit_keeps_shared_identity_and_frozen_snapshots(client):
    m = material(client); p = product(client, m)
    other = product(client, m, code='OTHER')
    o = order(client); sup = supply(client, m)
    edited = command(client, 'product_save', p['data'] | {'materials': [
        dict(material_id=m['id'], material_version=m['version'], material=m['data'] | {'supplier': '更新供应商', 'net_weight_kg': '.009'}, usage='2')
    ]}, p)
    rows = workspace(client)['records']
    current = next(r for r in rows if r['id'] == m['id'])
    assert current['version'] == m['version'] + 1 and current['data']['supplier'] == '更新供应商'
    assert edited['data']['materials'] == [dict(material_id=m['id'], usage='2')]
    assert next(r for r in rows if r['id'] == other['id']) == other
    assert next(r for r in rows if r['id'] == o['id']) == o
    assert next(r for r in rows if r['id'] == sup['id']) == sup


@pytest.mark.parametrize('change', [{'unit': 'SET'}, {'specification': 'OTHER'}, {'code': 'OTHER'}])
def test_product_nested_material_edit_preserves_used_identity_rules(client, change):
    m = material(client); p = product(client, m); order(client)
    before = workspace(client)['records']
    command(client, 'product_save', p['data'] | {'materials': [
        dict(material_id=m['id'], material_version=m['version'], material=m['data'] | change, usage='1')
    ]}, p, status=422)
    assert workspace(client)['records'] == before


def test_product_nested_material_cannot_cross_factory_or_bypass_permissions(client):
    customer(client); m = material(client)
    before = workspace(client)['records']
    client.materials_auth['user'] = actor('read', 'write', factory='huadeng')
    command(client, 'customer_save', dict(name='合成客户'), factory='huadeng')
    command(client, 'product_save', dict(code='CROSS', customer='合成客户', materials=[
        dict(material=inline_material('ORPHAN'), usage='1'),
        dict(material_id=m['id'], material_version=m['version'], material=m['data'], usage='1')
    ]), factory='huadeng', status=404)
    assert client.get(BASE + '/workspace', params={'factory_id': 'huadeng'}).json()['records'][0]['kind'] == 'customer'
    assert len(client.get(BASE + '/workspace', params={'factory_id': 'huadeng'}).json()['records']) == 1
    client.materials_auth['user'] = actor('read')
    command(client, 'product_save', dict(code='DENIED', customer='合成客户', materials=[dict(material=inline_material('DENIED'), usage='1')]), status=403)
    assert workspace(client)['records'] == before


@pytest.mark.parametrize('item', [dict(usage='1'), dict(material=inline_material('NEW'), material_version=3, usage='1'),
                                  dict(material_id='existing', material=inline_material('EDIT'), usage='1')])
def test_product_nested_material_requires_a_source_and_original_edit_revision(client, item):
    customer(client)
    command(client, 'product_save', dict(code='INVALID', customer='合成客户', materials=[item]), status=422)
    assert len(workspace(client)['records']) == 1


def product(client, m, code='ITEM01', **changes):
    ensure_customer(client, changes.get('customer', '合成客户'))
    return command(client, 'product_save', dict(code=code, customer='合成客户', daily_capacity='10',
                    materials=[dict(material_id=m['id'], usage='1')]) | changes)


def order(client, m=None, contract='TEST-A', quantity='100', **changes):
    return command(client, 'order_save', dict(customer='合成客户', recipient='RRI', contract_no=contract,
                    item_no='ITEM01', batch='1', quantity=quantity, delivery_date='2026-11-30') | changes)


def supply(client, m, normal='90', spare='10', **changes):
    return command(client, 'supply_save', dict(material_id=m['id'], purchase_no='PO-TEST', purchase_line='0001',
                    purchase_date='2026-10-01', supplier='合成供应商', purchase_quantity='100', normal_quantity=normal,
                    spare_quantity=spare, batch='SUPPLY-1', price='.10', currency='USD', billing_weight_source='net_weight_kg', promised_date='2026-10-21') | changes)


def allocation(client, sup, ord, normal='90', spare='10', **changes):
    return command(client, 'allocate', dict(supply_id=sup['id'], order_id=ord['id'], demand_id=ord['data']['demands'][0]['id'],
                    normal_quantity=normal, spare_quantity=spare) | changes)


def historical_allocation(client, sup, ord, normal='90', spare='10'):
    """Seed pre-policy history, never bypass receipt validation in current commands."""
    with Session(client.engine) as db:
        row = service.Ledger(db, 'huaxing').save('allocation', dict(supply_id=sup['id'], order_id=ord['id'],
            demand_id=ord['data']['demands'][0]['id'], normal_quantity=normal, spare_quantity=spare, reason='历史分配'))
        result = service.view(row); db.commit()
        return result


def received_flow(client, quantity='100'):
    m = material(client); p = product(client, m); o = order(client, quantity=quantity)
    sup = supply(client, m)
    sup = command(client, 'receipt_confirm', {'actual_date': '2026-10-22'}, sup)
    return m, p, o, sup


def shipment(client, alloc, normal='9', spare='1', **changes):
    item = dict(allocation_id=alloc['id'], normal_quantity=normal, spare_quantity=spare)
    package = dict(label='合成测试箱', cartons=1, length_cm='27', width_cm='14.5', height_cm='27.3',
                   carton_weight_kg='.020', items=[item])
    return command(client, 'shipment_save', dict(customer='合成客户', recipient='RRI', shipment_date='2026-10-23',
                                                lines=[item], packages=[package]) | changes)


def workspace(client):
    response = client.get(BASE + '/workspace', params={'factory_id': 'huaxing'})
    assert response.status_code == 200, response.text
    return response.json()


def test_production_workdays_and_natural_days():
    plan = production_plan('100', '10', '2026-11-30', 25)
    assert plan == dict(production_days=10, start_date='2026-11-16', dispatch_date='2026-10-22',
                        required_date='2026-10-22', transport_days=25, preparation_days=0)
    assert production_plan('101', '10', '2026-11-30', 25)['production_days'] == 11
    assert production_plan('100', None, '2026-11-30', 25) is None


def test_empty_real_workspace_and_factory_scope(client):
    assert workspace(client)['records'] == []
    assert client.get(BASE + '/workspace?factory_id=group').status_code == 422
    assert client.get(BASE + '/workspace?factory_id=huakang-a').status_code == 403
    assert client.get(BASE + '/workspace?factory_id=huaxing&factory_id=huakang-a').status_code == 422
    m = material(client)
    client.materials_auth['user'] = actor('read', 'write', factory='huakang-a')
    assert command(client, 'material_save', m['data'], m, factory='huakang-a', status=404).status_code == 404


def test_manual_demand_and_pinned_master_recalculation(client):
    m = material(client); p = product(client, m)
    o = order(client, quantity='100')
    d = o['data']['demands'][0]
    assert d['computed_quantity'] == '100'
    assert d['purchase_start'] == '2026-10-04'
    o = command(client, 'demand_adjust', dict(demand_id=d['id'], quantity='120.12345678', reason='合成调整'), o)
    p = command(client, 'product_save', p['data'] | {'materials': [{'material_id': m['id'], 'usage': '2'}]}, p)
    assert workspace(client)['board'][0]['demands'][0]['computed_quantity'] == '100'
    o = command(client, 'order_recalculate', {'reason': '合成显式重算'}, o)
    assert o['data']['demands'][0]['computed_quantity'] == '200'
    assert o['data']['demands'][0]['quantity'] == '120.12345678'
    assert o['data']['demands'][0]['id'] == d['id']


def test_fixed_customer_identity_scope_versions_and_legacy_contract_edits(client):
    c = customer(client, '  合成客户  ')
    assert c['kind'] == 'customer' and c['data']['name'] == '合成客户'
    command(client, 'customer_save', c['data'], status=409)
    command(client, 'customer_save', c['data'] | {'name': '覆盖关联身份'}, c, status=422)
    m = material(client)
    command(client, 'product_save', dict(code='X', customer='未配置客户', materials=[]), status=422)
    p = product(client, m)
    assert p['data']['customer_id'] == c['id']
    old = order(client)
    before = deepcopy(old['data']['customer_snapshot'])
    client.materials_auth['user'] = actor('read', 'write', factory='huakang-a')
    other = command(client, 'customer_save', c['data'], factory='huakang-a')
    command(client, 'customer_save', c['data'], c, factory='huakang-a', status=404)
    assert other['id'] != c['id']
    client.materials_auth['user'] = actor('read', 'write', 'receipt_confirm', 'export')
    assert all(r['id'] != other['id'] for r in workspace(client)['records'])
    command(client, 'product_save', p['data'] | {'customer_id': other['id']}, p, status=422)
    inactive = command(client, 'customer_save', c['data'] | {'status': 'inactive'}, c)
    command(client, 'customer_save', c['data'], c, status=409)
    command(client, 'order_save', dict(customer='合成客户', recipient='RRI', contract_no='NEW', item_no='ITEM01', batch='1', quantity='1', delivery_date='2026-11-30'), status=422)
    command(client, 'order_recalculate', {'reason': '禁止停用客户重新生成'}, old, status=422)
    changed = command(client, 'order_save', {key: old['data'][key] for key in ('customer', 'recipient', 'contract_no', 'po_no', 'item_no', 'batch', 'quantity', 'unit', 'delivery_date')} | {'reason': '历史合同核对', 'quantity': '101'}, old)
    assert changed['data']['customer_snapshot'] == before
    client.materials_auth['user'] = actor('read', 'receipt_confirm', department='pmc-warehouse')
    command(client, 'customer_save', inactive['data'], inactive, status=403)


def test_customer_product_and_material_recipe_are_required_for_new_contracts(client):
    data = dict(customer='合成客户', recipient='RRI', contract_no='NEW', item_no='ITEM01', batch='1', quantity='3', delivery_date='2026-11-30')
    command(client, 'order_save', data, status=422)
    customer(client)
    command(client, 'order_save', data, status=422)
    p = command(client, 'product_save', dict(code='ITEM01', customer='合成客户', materials=[]))
    command(client, 'order_save', data, status=422)
    m = material(client)
    p = command(client, 'product_save', p['data'] | {'materials': [dict(material_id=m['id'], usage='4')]}, p)
    created = command(client, 'order_save', data)
    assert created['data']['demands'][0]['quantity'] == '12'
    customer(client, '另一客户')
    command(client, 'order_save', data | {'customer': '另一客户'}, status=422)


def test_import_catalog_errors_recover_and_confirmation_rechecks_customer_status(client):
    c = customer(client)
    payload = dict(factory_id='huaxing', customer='合成客户', recipient='RRI')
    content = schedule_bytes()
    def preview():
        result = client.post(BASE + '/imports/preview', data=payload, files={'file': ('schedule.xlsx', content)})
        assert result.status_code == 200, result.text
        return result.json()['record']
    first = preview()
    assert first['data']['rows'][0]['blocked']
    product(client, material(client), code='00001')
    recovered = next(r for r in workspace(client)['records'] if r['id'] == first['id'])
    assert recovered['id'] == first['id'] and not recovered['data']['rows'][0]['blocked']
    assert preview() == recovered
    with Session(client.engine) as db:
        assert db.get(Record, first['id']).data['rows'][0]['blocked']  # Reading must not rewrite source evidence.
    inactive = command(client, 'customer_save', c['data'] | {'status': 'inactive'}, c)
    command(client, 'import_confirm', {'indexes': [0]}, recovered, status=422)
    assert not any(r['kind'] == 'order' for r in workspace(client)['records'])
    command(client, 'customer_save', inactive['data'] | {'status': 'active'}, inactive)
    confirmed = command(client, 'import_confirm', {'indexes': [0]}, recovered)
    assert confirmed['data']['status'] == 'confirmed'


def test_missing_capacity_does_not_hide_actual_complete(client):
    m = material(client, procurement_days=None); product(client, m, daily_capacity=None)
    o = order(client); sup = supply(client, m)
    assert workspace(client)['board'][0]['status'] == 'missing_parameters'
    sup = command(client, 'receipt_confirm', {'actual_date': '2026-10-22'}, sup)
    allocation(client, sup, o)
    board = workspace(client)['board'][0]
    assert board['status'] == 'complete' and board['planning_status'] == 'missing_parameters'


def test_reply_does_not_receive_or_hide_late_risk(client):
    m = material(client); product(client, m); o = order(client); sup = supply(client, m)
    historical_allocation(client, sup, o)
    sup = command(client, 'supply_reply', {'promised_date': '2026-10-25', 'reason': '来料仓合成回复'}, sup)
    row = workspace(client)['board'][0]
    assert row['status'] == 'shortage' and row['demands'][0]['coverage'] == '0'
    assert row['demands'][0]['risk'] == 'late' and row['plan']['required_date'] == '2026-10-22'
    assert sup['data']['status'] == 'expected' and sup['data']['actual_date'] is None


def test_receipt_independent_permission_and_explicit_deny(client):
    m = material(client); sup = supply(client, m)
    client.materials_auth['user'] = actor('read', 'write')
    command(client, 'receipt_confirm', {'actual_date': '2026-10-22'}, sup, status=403)
    client.materials_auth['user'] = actor('read', 'receipt_confirm', department='pmc-warehouse')
    command(client, 'receipt_confirm', {'actual_date': '2026-10-22'}, sup)
    command(client, 'material_save', m['data'], m, status=403)
    deny = AuthOverrideContext(id='test-deny', permission_code='indonesia_material:read', factory_id='huaxing', department='*', effect='deny')
    client.materials_auth['user'] = replace(actor('read', 'write'), overrides=(deny,))
    assert client.get(BASE + '/workspace?factory_id=huaxing').status_code == 403


def test_purchase_cap_includes_free_spares_and_all_batches(client):
    m = material(client); supply(client, m)
    command(client, 'supply_save', dict(material_id=m['id'], purchase_no='PO-TEST', purchase_line='0001', purchase_date='2026-10-01',
                    supplier='合成供应商', purchase_quantity='100', normal_quantity='1', spare_quantity='0', batch='SUPPLY-2', price='.1'), status=409)
    assert len([r for r in workspace(client)['records'] if r['kind'] == 'supply']) == 1


def test_allocations_do_not_double_reserve_and_failed_reallocate_rolls_back(client):
    m, _, a, sup = received_flow(client, quantity='50')
    b = order(client, contract='TEST-B', quantity='50')
    alloc_a = allocation(client, sup, a, normal='50', spare='0')
    allocation(client, sup, b, normal='40', spare='10')
    command(client, 'allocate', dict(supply_id=sup['id'], order_id=a['id'], demand_id=a['data']['demands'][0]['id'],
                                    normal_quantity='1', spare_quantity='0'), status=409)
    command(client, 'reallocate', dict(order_id=b['id'], demand_id=b['data']['demands'][0]['id'], normal_quantity='10', spare_quantity='0', reason='测试超配'), alloc_a, status=409)
    original = next(r for r in workspace(client)['records'] if r['id'] == alloc_a['id'])
    assert original['version'] == 1 and original['data']['normal_quantity'] == '50'


def test_reallocate_unshipped_reopens_original_shortage(client):
    m, _, a, sup = received_flow(client, quantity='50')
    b = order(client, contract='TEST-B', quantity='50')
    alloc = allocation(client, sup, a, normal='50', spare='0')
    command(client, 'reallocate', dict(order_id=b['id'], demand_id=b['data']['demands'][0]['id'], normal_quantity='10', spare_quantity='0', reason='测试调整'), alloc)
    board = {r['contract_no']: r for r in workspace(client)['board']}
    assert board['TEST-A']['demands'][0]['shortage'] == '10'
    assert board['TEST-B']['demands'][0]['coverage'] == '10'


def test_version_conflicts_and_exact_request_replay(client):
    m = material(client)
    updated = command(client, 'material_save', m['data'] | {'name': '合成新版'}, m)
    command(client, 'material_save', m['data'], m, status=409)
    op = uuid4().hex
    first = command(client, 'material_save', updated['data'] | {'name': '合成再次修订'}, updated, operation_id=op)
    replay = command(client, 'material_save', updated['data'] | {'name': '合成再次修订'}, updated, operation_id=op)
    assert replay == first
    command(client, 'material_save', updated['data'] | {'name': '错误复用'}, updated, operation_id=op, status=409)


def test_confirm_export_is_immutable_and_normal_shipment_preserves_fulfilment(client):
    _, _, o, sup = received_flow(client)
    alloc = allocation(client, sup, o)
    draft = shipment(client, alloc)
    op = uuid4().hex
    confirmed = command(client, 'shipment_confirm', {}, draft, operation_id=op)
    assert command(client, 'shipment_confirm', {}, draft, operation_id=op) == confirmed
    assert workspace(client)['board'][0]['status'] == 'complete'
    assert confirmed['data']['packing_totals']['lines'][alloc['id']]['gross_weight_kg'] == '.030' or Decimal(confirmed['data']['packing_totals']['lines'][alloc['id']]['gross_weight_kg']) == Decimal('.030')
    file_a = client.get(BASE + f'/shipments/{confirmed["id"]}/export?factory_id=huaxing')
    file_b = client.get(BASE + f'/shipments/{confirmed["id"]}/export?factory_id=huaxing')
    assert file_a.content == file_b.content and file_a.status_code == 200
    record = next(r for r in workspace(client)['records'] if r['id'] == alloc['id'])
    assert record['data']['shipped_normal'] == '9' and record['data']['shipped_spare'] == '1'
    with Session(client.engine) as db:
        assert db.scalar(select(func.count()).select_from(Asset)) == 1
    wb = detail_book(BytesIO(file_a.content), data_only=True)
    sheet = wb[SHEET_NAME]
    assert sheet['L4'].value == 9 and sheet['K4'].value == .01
    assert sheet['AP4'].value == .9 and sheet['P4'].value == .03
    summary = confirmed['data']['supplier_summary']
    assert len(summary) == 1 and summary[0]['materials'] == 1
    assert summary[0]['amount'] == '0.90'
    assert next(r for r in workspace(client)['records'] if r['id'] == confirmed['id'])['data']['supplier_summary'] == summary
    assert all(sheet[f'{col}4'].value is None for col in ('B', 'C', 'E', 'N', 'O', 'V', 'Z', 'AA', 'AG'))
    assert sheet['AS4'].value == '華登製品實業有限公司'
    assert '备品 1 PCS，免费' in sheet['BE4'].value
    assert wb[CALC_NAME].sheet_state == 'hidden'
    # The latest request removes tutorial images; only this shipment's actual photo survives.
    with ZipFile(BytesIO(file_a.content)) as exported, ZipFile(TEMPLATE) as template:
        original_styles = ET.fromstring(template.read('xl/styles.xml'))
        output_styles = ET.fromstring(exported.read('xl/styles.xml'))
        for section in ('fonts', 'fills', 'borders', 'numFmts'):
            assert ET.tostring(original_styles.find('s:'+section, NS)) == ET.tostring(output_styles.find('s:'+section, NS))
        drawing = ET.fromstring(exported.read(DRAWING_PART))
        assert len(drawing.findall('.//xdr:pic', NS)) == 1
        assert len([n for n in exported.namelist() if n.startswith('xl/media/indonesia-material-')]) == 1
        assert all(exported.read(n) == template.read(n) for n in ('xl/worksheets/sheet1.xml', 'xl/worksheets/sheet3.xml'))
    client.materials_auth['user'] = actor('read')
    assert client.get(BASE + f'/shipments/{confirmed["id"]}/export?factory_id=huaxing').status_code == 403
    assert client.get(BASE + f'/assets/{confirmed["data"]["export_asset_id"]}?factory_id=huaxing').status_code == 403


def test_draft_packages_require_separate_normal_and_spare_reconciliation(client):
    _, _, o, sup = received_flow(client); alloc = allocation(client, sup, o)
    draft = shipment(client, alloc)
    changed = deepcopy(draft['data']); changed.pop('status'); changed['packages'][0]['items'][0].update(normal_quantity='10', spare_quantity='0')
    draft = command(client, 'shipment_save', changed, draft)
    command(client, 'shipment_confirm', {}, draft, status=422)
    assert next(r for r in workspace(client)['records'] if r['id'] == draft['id'])['data']['status'] == 'draft'


def test_mixed_carton_can_confirm_once_with_frozen_material_rows(client):
    m = material(client); other = material(client, code='00002')
    product(client, m, materials=[dict(material_id=x['id'], usage='1') for x in (m, other)])
    o = order(client)
    selected = []
    for i, entry in enumerate((m, other)):
        sup = supply(client, entry, purchase_line=str(i+1))
        sup = command(client, 'receipt_confirm', {'actual_date': '2026-10-22'}, sup)
        a = allocation(client, sup, o, demand_id=o['data']['demands'][i]['id'])
        selected.append(dict(allocation_id=a['id'], normal_quantity='3', spare_quantity='1'))
    draft = command(client, 'shipment_save', dict(customer='合成客户', recipient='RRI', shipment_date='2026-10-23',
                    lines=selected, packages=[dict(label='真实混箱', cartons=1, length_cm='27', width_cm='14.5',
                    height_cm='27.3', carton_weight_kg='.02', items=selected)]))
    confirmed = command(client, 'shipment_confirm', {}, draft)
    assert confirmed['data']['export_material_count'] == 2 and confirmed['data']['export_layout_version'] == 'material_row_v2'
    content = client.get(BASE + f'/shipments/{draft["id"]}/export?factory_id=huaxing').content
    sheet = detail_book(BytesIO(content), data_only=True)[SHEET_NAME]
    assert {str(r) for r in sheet.merged_cells.ranges} >= {'AT4:AT5', 'S4:S5'}
    assert sheet['AT4'].value == 1 and sheet['AT5'].value is None
    assert sheet['P4'].value == pytest.approx(.014) and sheet['P5'].value == pytest.approx(.014)


@pytest.mark.parametrize('conflict', ['dimensions', 'currency'])
def test_export_conflicts_leave_draft_and_allocation_unshipped(client, conflict):
    m, _, o, sup = received_flow(client); a = allocation(client, sup, o)
    selected = [dict(allocation_id=a['id'], normal_quantity='9', spare_quantity='1')]
    packages = [dict(label='整箱', cartons=1, length_cm='27', width_cm='14.5', height_cm='27.3',
                     carton_weight_kg='.02', items=deepcopy(selected))]
    if conflict == 'dimensions':
        packages[0]['items'] = [dict(allocation_id=a['id'], normal_quantity='5', spare_quantity='0')]
        packages.append(dict(label='尾箱', cartons=1, length_cm='28', width_cm='14.5', height_cm='27.3',
                             carton_weight_kg='.02', items=[dict(allocation_id=a['id'], normal_quantity='4', spare_quantity='1')]))
    else:
        other = material(client, code='00002'); product(client, other, code='ITEM02')
        other_order = order(client, contract='TEST-B', item_no='ITEM02')
        source = supply(client, other, purchase_line='2', currency='CNY')
        source = command(client, 'receipt_confirm', {'actual_date': '2026-10-22'}, source)
        other_a = allocation(client, source, other_order)
        selected.append(dict(allocation_id=other_a['id'], normal_quantity='9', spare_quantity='1'))
        packages.append(dict(label='第二物料', cartons=1, length_cm='27', width_cm='14.5', height_cm='27.3',
                             carton_weight_kg='.02', items=[selected[-1]]))
    draft = command(client, 'shipment_save', dict(customer='合成客户', recipient='RRI', shipment_date='2026-10-23',
                                                lines=selected, packages=packages))
    failure = command(client, 'shipment_confirm', {}, draft, status=422)
    assert ('尺寸' in failure.text or '长宽高' in failure.text) if conflict == 'dimensions' else '采购币种不一致' in failure.text
    records = {r['id']: r for r in workspace(client)['records']}
    assert records[draft['id']]['data']['status'] == 'draft'
    assert records[a['id']]['data']['shipped_normal'] == '0' and records[a['id']]['data']['shipped_spare'] == '0'
    with Session(client.engine) as db:
        assert db.scalar(select(func.count()).select_from(Asset)) == 0


@pytest.mark.parametrize('net', [None, '0', '0.00123456'])
def test_master_saves_derive_compatible_billing_weight_from_net_for_direct_and_nested_inputs(client, net):
    m = material(client, net_weight_kg=net, billing_weight_kg='999')
    assert m['data']['billing_weight_kg'] == m['data']['net_weight_kg']
    customer(client)
    p = command(client, 'product_save', dict(code='NET', customer='合成客户', materials=[
        dict(material=inline_material('INLINE-NET', net_weight_kg=net, billing_weight_kg='777'), usage='1')]))
    nested = next(r for r in workspace(client)['records'] if r['id'] == p['data']['materials'][0]['material_id'])
    assert nested['data']['billing_weight_kg'] == nested['data']['net_weight_kg']


def test_legacy_uncommitted_purchase_retries_keep_the_old_basis_and_new_commands_use_net(client):
    m = material(client); product(client, m); o = order(client)
    # Synthetic pre-upgrade master, with a genuinely different old billing weight.
    with Session(client.engine) as db:
        master = db.get(Record, m['id'])
        master.data = master.data | {'billing_weight_kg': '.002'}
        db.commit()
    legacy_payload = dict(material_id=m['id'], purchase_no='LEGACY-PO', purchase_line='1', purchase_date='2026-10-01',
                          supplier='合成供应商', purchase_quantity='100', normal_quantity='90', spare_quantity='10',
                          batch='OLD', price='100', currency='USD', billing_basis='weight', billing_factor='1')
    operation = uuid4().hex
    with Session(client.engine) as db:
        assert db.get(Receipt, operation) is None  # Durable old command never reached the server.
    old = command(client, 'supply_save', legacy_payload, operation_id=operation)
    assert 'billing_weight_source' not in old['data']
    assert old['data']['material_snapshot']['data']['billing_weight_kg'] == '.002'
    m = command(client, 'material_save', m['data'] | {'net_weight_kg': '.005'}, m)
    assert command(client, 'supply_save', legacy_payload, operation_id=operation) == old
    editable = {key: value for key, value in old['data'].items() if key in service.s.Supply.model_fields}
    old = command(client, 'supply_save', editable | {'normal_quantity': '89', 'billing_weight_source': 'net_weight_kg'}, old)
    assert 'billing_weight_source' not in old['data']
    assert old['data']['material_snapshot']['data']['billing_weight_kg'] == '.002'
    old = command(client, 'receipt_confirm', {'actual_date': '2026-10-22'}, old)
    old_allocation = allocation(client, old, o, normal='9', spare='1')
    sh = command(client, 'shipment_confirm', {}, shipment(client, old_allocation))
    assert Decimal(sh['data']['snapshot_lines'][0]['amount']) == Decimal('1.8')
    original_bytes = client.get(BASE + f"/shipments/{sh['id']}/export", params={'factory_id': 'huaxing'}).content
    new = supply(client, m, purchase_no='CURRENT-PO', billing_basis='weight', price='100')
    assert new['data']['billing_weight_source'] == 'net_weight_kg'
    assert new['data']['material_snapshot']['data']['billing_weight_kg'] == '0.005'
    new = command(client, 'receipt_confirm', {'actual_date': '2026-10-22'}, new)
    new_allocation = allocation(client, new, o, normal='9', spare='1')
    current_shipment = command(client, 'shipment_confirm', {}, shipment(client, new_allocation))
    assert Decimal(current_shipment['data']['snapshot_lines'][0]['amount']) == Decimal('4.5')
    command(client, 'material_save', m['data'] | {'net_weight_kg': '.009'}, m)
    assert client.get(BASE + f"/shipments/{sh['id']}/export", params={'factory_id': 'huaxing'}).content == original_bytes


def test_new_purchase_from_unedited_legacy_master_freezes_net_on_quantity_edit(client):
    m = material(client)
    with Session(client.engine) as db:
        master = db.get(Record, m['id'])
        master.data = master.data | {'billing_weight_kg': '999'}
        db.commit()
    sup = supply(client, m, billing_basis='weight')
    assert sup['data']['billing_weight_source'] == 'net_weight_kg'
    assert sup['data']['material_snapshot']['data']['billing_weight_kg'] == '0.001'
    command(client, 'material_save', m['data'] | {'net_weight_kg': '.005'}, m)
    editable = {key: value for key, value in sup['data'].items() if key in service.s.Supply.model_fields}
    changed = command(client, 'supply_save', editable | {'normal_quantity': '89', 'billing_weight_source': None}, sup)
    assert changed['data']['billing_weight_source'] == 'net_weight_kg'
    assert changed['data']['material_snapshot'] == sup['data']['material_snapshot']


@pytest.mark.parametrize('net', [None, '0'])
def test_current_weight_purchase_requires_net_even_if_legacy_billing_weight_exists(client, net):
    m = material(client, net_weight_kg=net)
    with Session(client.engine) as db:
        master = db.get(Record, m['id'])
        master.data = master.data | {'billing_weight_kg': '999'}
        db.commit()
    product(client, m); o = order(client)
    sup = supply(client, m, billing_basis='weight', price='100')
    sup = command(client, 'receipt_confirm', {'actual_date': '2026-10-22'}, sup)
    alloc = allocation(client, sup, o, normal='9', spare='1')
    draft = shipment(client, alloc)
    command(client, 'shipment_confirm', {}, draft, status=422)
    m = command(client, 'material_save', m['data'] | {'net_weight_kg': '.003', 'billing_weight_kg': '888'}, m)
    sup = command(client, 'output_material_refresh', {'reason': '补齐净重'}, sup)
    assert sup['data']['measured_weights']['billing_weight_kg'] == '0.003'
    confirmed = command(client, 'shipment_confirm', {}, draft)
    assert Decimal(confirmed['data']['snapshot_lines'][0]['amount']) == Decimal('2.7')


def test_weight_billing_uses_net_weight_and_never_charges_spares(client):
    m = material(client); product(client, m); o = order(client)
    sup = supply(client, m, billing_basis='weight', billing_factor='1', price='100')
    sup = command(client, 'receipt_confirm', {'actual_date': '2026-10-22'}, sup)
    alloc = allocation(client, sup, o)
    confirmed = command(client, 'shipment_confirm', {}, shipment(client, alloc))
    line = confirmed['data']['snapshot_lines'][0]
    assert Decimal(line['amount']) == Decimal('.9')  # 9 * .001 * 100, spare is free.
    assert Decimal(confirmed['data']['packing_totals']['lines'][alloc['id']]['net_weight_kg']) == Decimal('.01')


def test_mixed_carton_weight_ratio_and_cbm_precision():
    lines = {'a': dict(normal_quantity='6', spare_quantity='0', material={'net_weight_kg': '1'}),
             'b': dict(normal_quantity='4', spare_quantity='0', material={'net_weight_kg': '1'})}
    package = dict(label='混箱', cartons=1, length_cm='27', width_cm='14.5', height_cm='27.3', carton_weight_kg='1', pallet='999',
                   items=[dict(allocation_id=k, normal_quantity=q, spare_quantity='0') for k, q in [('a','6'),('b','4')]])
    calculated = packaging_totals([package], lines)
    assert Decimal(calculated['lines']['a']['gross_weight_kg']) == Decimal('6.6')
    assert Decimal(calculated['lines']['b']['gross_weight_kg']) == Decimal('4.4')
    package.update(cartons=53, carton_weight_kg='0')
    assert packaging_totals([package], lines)['cbm'] == '0.57'


def schedule_bytes(*, duplicate=False):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'ITEM表'
    ws.append([]); ws.append([])
    ws.append(['CONTRACT NO.', 'SO#/REFERENCE', 'P/O#:', '产品编号', '数量', '客要求走货期', '客名'])
    ws.append(['TEST-I', 'SO-TEST', 'PO-TEST', '00001', 100, '2026-11-30', '合成客户'])
    if duplicate:
        ws.append(['TEST-I', 'SO-TEST', 'PO-TEST', '00001', 50, '2026-12-15', '合成客户'])
    stream = BytesIO(); wb.save(stream); return stream.getvalue()


def test_import_selected_rows_repeated_files_and_original_lineage(client):
    product(client, material(client), code='00001')
    original = schedule_bytes()
    data = dict(factory_id='huaxing', customer='合成客户', recipient='RRI')
    result = client.post(BASE + '/imports/preview', data=data, files={'file': ('schedule.xlsx', original)})
    assert result.status_code == 200, result.text
    preview = result.json()['record']
    confirmed = command(client, 'import_confirm', {'indexes': [0]}, preview)
    repeated = client.post(BASE + '/imports/preview', data=data, files={'file': ('schedule.xlsx', original)})
    assert repeated.json()['record']['id'] == confirmed['id'] and repeated.json()['record']['data']['status'] == 'confirmed'
    rows = workspace(client)['records']; orders = [r for r in rows if r['kind'] == 'order']
    assert len(orders) == 1 and orders[0]['data']['item_no'] == '00001'
    assert orders[0]['data']['source']['sheet'] == 'ITEM表' and orders[0]['data']['source']['row'] == 4
    saved = client.get(BASE + f'/assets/{preview["data"]["asset_id"]}?factory_id=huaxing')
    assert saved.content == original


def test_ambiguous_delivery_rows_and_caches_are_blocked(client):
    customer(client)
    result = client.post(BASE + '/imports/preview', data=dict(factory_id='huaxing', customer='合成客户', recipient='RRI'),
                         files={'file': ('schedule.xlsx', schedule_bytes(duplicate=True))})
    assert result.status_code == 200, result.text
    record = result.json()['record']
    assert all(r['blocked'] for r in record['data']['rows'])
    command(client, 'import_confirm', {'indexes': [0]}, record, status=422)
    assert not any(r['kind'] == 'order' for r in workspace(client)['records'])


def test_referenced_material_specification_and_units_are_identity(client):
    m = material(client, specification='M2'); product(client, m); order(client)
    command(client, 'material_save', m['data'] | {'specification': 'M3'}, m, status=422)
    command(client, 'material_save', m['data'] | {'unit': 'KGM'}, m, status=422)
    updated = command(client, 'material_save', m['data'] | {'english_name': 'Corrected description'}, m)
    assert updated['data']['specification'] == 'M2'


def test_explicit_output_refresh_fills_missing_data_but_preserves_known_weights(client):
    m = material(client, english_name='', image_data='', net_weight_kg=None, billing_weight_kg=None)
    product(client, m); o = order(client); sup = supply(client, m, billing_basis='weight', price='100')
    sup = command(client, 'receipt_confirm', {'actual_date': '2026-10-22'}, sup)
    alloc = allocation(client, sup, o); draft = shipment(client, alloc)
    command(client, 'shipment_confirm', {}, draft, status=422)
    image = BytesIO(); Image.new('RGB', (4, 4), 'blue').save(image, 'PNG')
    m = command(client, 'material_save', m['data'] | dict(english_name='Completed description', net_weight_kg='.001',
        billing_weight_kg='.002', image_data='data:image/png;base64,' + base64.b64encode(image.getvalue()).decode()), m)
    sup = command(client, 'output_material_refresh', {'reason': '补齐缺失资料'}, sup)
    first = command(client, 'shipment_confirm', {}, draft)
    assert Decimal(first['data']['snapshot_lines'][0]['amount']) == Decimal('.9')
    # Subsequent master corrections do not change known source measurements or an existing file.
    m = command(client, 'material_save', m['data'] | {'net_weight_kg': '.5', 'billing_weight_kg': '.7'}, m)
    sup = command(client, 'output_material_refresh', {'reason': '更新文字和图片版本'}, sup)
    second = command(client, 'shipment_confirm', {}, shipment(client, alloc, normal='1', spare='0'))
    # First explicit missing-measurement completion becomes this batch's fixed basis.
    assert Decimal(second['data']['snapshot_lines'][0]['material']['net_weight_kg']) == Decimal('.001')


def test_received_board_preserves_promised_and_actual_dates(client):
    m = material(client); product(client, m); o = order(client); sup = supply(client, m, promised_date='2026-10-25')
    sup = command(client, 'receipt_confirm', {'actual_date': '2026-10-26'}, sup)
    allocation(client, sup, o)
    demand = workspace(client)['board'][0]['demands'][0]
    assert demand['promised_date'] == '2026-10-25' and demand['actual_date'] == '2026-10-26'
    assert demand['risk'] == 'late_received'


def test_reimport_previews_changes_and_allows_new_rows_independently(client):
    m = material(client)
    product(client, m, code='00001')
    product(client, m, code='00002')
    order(client, contract='TEST-I', item_no='00001', po_no='PO-TEST')
    wb = openpyxl.load_workbook(BytesIO(schedule_bytes())); sheet = wb['ITEM表']
    sheet['F4'] = '2026-12-15'
    sheet.append(['TEST-NEW', 'SO-NEW', 'PO-NEW', '00002', 50, '2026-12-15', '合成客户'])
    content = BytesIO(); wb.save(content)
    response = client.post(BASE + '/imports/preview', data=dict(factory_id='huaxing', customer='合成客户', recipient='RRI'),
                           files={'file': ('changed.xlsx', content.getvalue())})
    record = response.json()['record']
    assert record['data']['rows'][0]['match_status'] == 'changed' and record['data']['rows'][0]['blocked']
    assert record['data']['rows'][0]['changes'][0]['field'] == 'delivery_date'
    assert record['data']['rows'][1]['match_status'] == 'new'
    command(client, 'import_confirm', {'indexes': [1]}, record)
    assert len(workspace(client)['board']) == 2


def test_partial_selected_import_can_be_completed_later(client):
    m = material(client)
    product(client, m, code='00001')
    product(client, m, code='00002')
    wb = openpyxl.load_workbook(BytesIO(schedule_bytes())); wb['ITEM表'].append(['TEST-B', 'SO-B', 'PO-B', '00002', 50, '2026-12-15', '合成客户'])
    content = BytesIO(); wb.save(content)
    response = client.post(BASE + '/imports/preview', data=dict(factory_id='huaxing', customer='合成客户', recipient='RRI'),
                           files={'file': ('partial.xlsx', content.getvalue())})
    record = response.json()['record']
    partial = command(client, 'import_confirm', {'indexes': [0]}, record)
    assert partial['data']['status'] == 'partial'
    assert partial['data']['rows'][0]['imported_order_id']
    confirmed = command(client, 'import_confirm', {'indexes': [1]}, partial)
    assert confirmed['data']['status'] == 'confirmed' and len(confirmed['data']['order_ids']) == 2


def test_file_job_return_rechecks_current_authorization(client, monkeypatch):
    async def revoke_during_parse(function, *args, db=None, **kwargs):
        if db:
            db.rollback()
        result = function(*args, **kwargs)
        client.materials_auth['user'] = actor('read')
        return result
    monkeypatch.setattr(api, 'run_order_job', revoke_during_parse)
    response = client.post(BASE + '/imports/preview', data=dict(factory_id='huaxing', customer='合成客户', recipient='RRI'),
                           files={'file': ('revoked.xlsx', schedule_bytes())})
    assert response.status_code == 403
    with Session(client.engine) as db:
        assert db.scalar(select(func.count()).select_from(Asset)) == 0
        assert db.scalar(select(func.count()).select_from(Record)) == 0


def test_concurrent_confirmations_never_over_ship(client):
    from concurrent.futures import ThreadPoolExecutor
    _, _, o, sup = received_flow(client); alloc = allocation(client, sup, o)
    drafts = [shipment(client, alloc, normal='60', spare='0') for _ in range(2)]
    bodies = [dict(factory_id='huaxing', operation_id=uuid4().hex, action='shipment_confirm', entity_id=d['id'],
                   expected_version=d['version'], data={}) for d in drafts]
    with ThreadPoolExecutor(max_workers=2) as executor:
        responses = list(executor.map(lambda body: client.post(BASE + '/commands', json=body), bodies))
    assert sorted(r.status_code for r in responses) == [200, 409]
    record = next(r for r in workspace(client)['records'] if r['id'] == alloc['id'])
    assert record['data']['shipped_normal'] == '60'
    with Session(client.engine) as db:
        assert db.scalar(select(func.count()).select_from(Asset)) == 1


def test_template_three_rows_cleans_guide_and_keeps_actual_photo_locations(client):
    _, _, o, sup = received_flow(client); alloc = allocation(client, sup, o)
    draft = shipment(client, alloc)
    data = deepcopy(draft['data']); data.pop('status')
    original_group = data['packages'][0]
    allocations = [alloc]
    for i in (2, 3):
        m = material(client, code=f'0000{i}')
        product(client, m, code=f'ITEM0{i}')
        other_order = order(client, contract=f'TEST-{i}', item_no=f'ITEM0{i}')
        source = supply(client, m, purchase_line=str(i))
        source = command(client, 'receipt_confirm', {'actual_date': '2026-10-22'}, source)
        allocations.append(allocation(client, source, other_order))
    data['lines'] = [dict(allocation_id=a['id'], normal_quantity='3', spare_quantity='0') for a in allocations]
    data['packages'] = []
    for i, a in enumerate(allocations):
        group = deepcopy(original_group); group['label'] = f'分组 {i+1}'
        group['items'] = [dict(allocation_id=a['id'], normal_quantity='3', spare_quantity='0')]
        data['packages'].append(group)
    draft = command(client, 'shipment_save', data, draft)
    confirmed = command(client, 'shipment_confirm', {}, draft)
    response = client.get(BASE + f'/shipments/{confirmed["id"]}/export?factory_id=huaxing')
    formulas = detail_book(BytesIO(response.content), data_only=False)
    sheet = formulas[SHEET_NAME]
    assert sheet['R6'].value == '=AW6*AX6*AY6/28316.75*0.0283'
    assert sheet['AQ4'].value.startswith("=SUM('出货计算'!J")
    cached = detail_book(BytesIO(response.content), data_only=True)[SHEET_NAME]
    assert cached['AQ4'].value == pytest.approx(.9)
    assert sheet.max_row == 6
    assert str(sheet.print_area) == f"'{SHEET_NAME}'!$A$3:$BE$6"
    with ZipFile(BytesIO(response.content)) as output:
        result = ET.fromstring(output.read(SHEET_PART))
        assert result.find("s:sheetData/s:row[@r='8']", NS) is None
        assert 'xl/comments2.xml' not in output.namelist()
        assert 'xl/drawings/vmlDrawing2.vml' not in output.namelist()
        picture = ET.fromstring(output.read(DRAWING_PART)).findall('.//xdr:oneCellAnchor',NS)
        assert [p.findtext('xdr:from/xdr:row',namespaces=NS) for p in picture] == ['3','4','5']
        assert all(p.findtext('xdr:from/xdr:col',namespaces=NS) == '19' for p in picture)


def test_migration_is_explicit_and_evidence_is_immutable(tmp_path):
    import importlib.util
    from pathlib import Path
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from sqlalchemy import text
    from sqlalchemy.exc import IntegrityError
    path = Path(__file__).parents[1] / 'alembic/versions/20261013_0160_indonesia_materials.py'
    spec = importlib.util.spec_from_file_location('materials_migration', path)
    migration = importlib.util.module_from_spec(spec); spec.loader.exec_module(migration)
    engine = create_engine('sqlite:///' + (tmp_path/'migration.db').as_posix())
    with engine.begin() as connection:
        with Operations.context(MigrationContext.configure(connection)):
            migration.upgrade()
        with pytest.raises(HTTPException, match='基础选项'):
            service.ready(Session(bind=connection))
        customer_path = path.with_name('20261021_0169_indonesia_customers.py')
        customer_spec = importlib.util.spec_from_file_location('materials_customer_migration', customer_path)
        customer_migration = importlib.util.module_from_spec(customer_spec); customer_spec.loader.exec_module(customer_migration)
        with Operations.context(MigrationContext.configure(connection)):
            customer_migration.upgrade()
        with pytest.raises(HTTPException, match='基础选项'):
            service.ready(Session(bind=connection))
        options_path = path.with_name('20261022_0173_indonesia_options.py')
        options_spec = importlib.util.spec_from_file_location('materials_options_migration', options_path)
        options_migration = importlib.util.module_from_spec(options_spec); options_spec.loader.exec_module(options_migration)
        with Operations.context(MigrationContext.configure(connection)):
            options_migration.upgrade()
        import runpy
        shared_migration = runpy.run_path(str(path.with_name('20261022_0174_indonesia_shared_options.py')))
        with Operations.context(MigrationContext.configure(connection)):
            shared_migration['upgrade']()
        assert service.ready(Session(bind=connection)) is None
        connection.execute(text("INSERT INTO indonesia_material_customers VALUES ('customer','huaxing','customer-key',1,'{}','2026-10-21')"))
        with Operations.context(MigrationContext.configure(connection)):
            with pytest.raises(RuntimeError, match='拒绝'):
                customer_migration.downgrade()
        connection.execute(text("INSERT INTO indonesia_material_assets VALUES ('a','huaxing','synthetic.xlsx','test','digest',X'01','2026-10-13')"))
        with pytest.raises(IntegrityError, match='immutable'):
            connection.execute(text("DELETE FROM indonesia_material_assets WHERE id='a'"))
        with Operations.context(MigrationContext.configure(connection)):
            with pytest.raises(RuntimeError, match='拒绝'):
                migration.downgrade()
    engine.dispose()


def test_permission_catalog_does_not_expand_existing_positions():
    from app.services.system_positions import SYSTEM_POSITION_DEFINITIONS, GENERAL_MANAGER_EXCLUDED_BUSINESS_PERMISSION_CODES
    from app.services.iam_scope import READ_PERMISSION_CODES
    assert 'indonesia_material:read' in READ_PERMISSION_CODES
    assert set(INDONESIA_MATERIAL_PERMISSION_CODES) <= GENERAL_MANAGER_EXCLUDED_BUSINESS_PERMISSION_CODES
    assert all(not (set(p.permission_codes) & set(INDONESIA_MATERIAL_PERMISSION_CODES)) for p in SYSTEM_POSITION_DEFINITIONS)


@pytest.mark.parametrize('change,status', [('revoke', 403), ('identity', 401), ('session', 401)])
def test_command_rechecks_identity_after_waiting_for_locks(client, monkeypatch, change, status):
    saved = material(client)
    original_lock = service.lock_transaction
    def changed_while_waiting(db, namespace, factory):
        original_lock(db, namespace, factory)
        if change == 'revoke':
            client.materials_auth['user'] = actor('read')
        elif change == 'identity':
            client.materials_auth['user'] = replace(client.materials_auth['user'], id='other-user')
        else:
            def expired(request, db):
                raise HTTPException(401, '登录已失效')
            monkeypatch.setattr(api, 'refresh_user', expired)
    monkeypatch.setattr(service, 'lock_transaction', changed_while_waiting)
    data = {key: value for key, value in saved['data'].items() if key in api.service.s.Material.model_fields}
    data['name'] = '不得保存的变更'
    command(client, 'material_save', data, saved, status=status)
    with Session(client.engine) as db:
        assert db.get(Record, saved['id']).data['name'] == saved['data']['name']
        assert db.scalar(select(func.count()).select_from(Receipt)) == 1


def test_receipt_rechecks_revocation_and_replay_never_bypasses_it(client, monkeypatch):
    saved = material(client)
    batch = supply(client, saved)
    operation = uuid4().hex
    body = dict(factory_id='huaxing', operation_id=operation, action='receipt_confirm',
                entity_id=batch['id'], expected_version=batch['version'], data={'actual_date':'2026-10-22'})
    assert client.post(BASE+'/commands', json=body).status_code == 200
    original_lock = service.lock_transaction
    def revoke_after_lock(db, namespace, factory):
        original_lock(db, namespace, factory)
        client.materials_auth['user'] = actor('read', 'write')
    monkeypatch.setattr(service, 'lock_transaction', revoke_after_lock)
    assert client.post(BASE+'/commands', json=body).status_code == 403
    with Session(client.engine) as db:
        assert db.scalar(select(func.count()).select_from(Receipt)) == 3


def test_fully_reallocated_old_supply_does_not_pollute_current_dates(client):
    m = material(client); product(client, m)
    a = order(client, contract='TEST-A', quantity='50'); b = order(client, contract='TEST-B', quantity='50')
    old = supply(client, m, normal='50', spare='0', promised_date='2026-10-25')
    old = command(client, 'receipt_confirm', {'actual_date': '2026-10-26'}, old)
    alloc = allocation(client, old, a, normal='50', spare='0')
    command(client, 'reallocate', dict(order_id=b['id'], demand_id=b['data']['demands'][0]['id'], normal_quantity='50', spare_quantity='0', reason='全部改绑'), alloc)
    new = supply(client, m, normal='50', spare='0', batch='SUPPLY-2', promised_date='2026-10-21')
    new = command(client, 'receipt_confirm', {'actual_date': '2026-10-22'}, new)
    allocation(client, new, a, normal='50', spare='0')
    current = next(r for r in workspace(client)['board'] if r['id'] == a['id'])['demands'][0]
    assert current['risk'] == 'received' and current['promised_date'] == '2026-10-21' and current['actual_date'] == '2026-10-22'
