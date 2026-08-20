from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.services.internal_quote_excel import _build_summary_sheet


def _summary_sheet(
    *,
    rate_percent: str | None,
    material_amount_hkd: str,
    misc_ratio: str = "0.02",
    region_code: str = "",
    indonesia_freight_hkd: str = "0",
    additional_tax_hkd: str = "0",
):
    workbook = Workbook()
    quote = SimpleNamespace(
        product_name="车缝退税测试产品",
        quote_no="IQ-SEWING-TAX",
        region_code=region_code,
        remark="",
    )
    rr2_cost_summary = {
        "t1": [{"key": "sewing_cloth", "label": "车衣", "value": "12.5000"}],
        "t2": [],
        "t3": [],
        "t4": [
            {
                "key": "sewcloth13",
                "label": "车缝物料退税",
                "amount_hkd": material_amount_hkd,
                "rate_percent": rate_percent,
                "deduction_hkd": None,
            }
        ],
        "shipping_pricing": {
            "markup": "1.2000",
            "active_markup_moq": "10000",
            "markup_tiers": [],
            "misc_ratio": misc_ratio,
            "additional_tax_hkd": additional_tax_hkd,
            "rows": [],
        },
    }
    # Keep the frozen snapshot at the historical 11.5% so the tests prove that
    # export follows the factory-aware rr2 row instead of falling back to it.
    reference_snapshot = {
        "fx": {"rmb_hkd": "0.85", "hkd_usd": "7.8"},
        "tax_rates": {"sewing_clothes": "0.115"},
    }
    _build_summary_sheet(
        workbook,
        quote,
        [],
        reference_snapshot,
        rr2_cost_summary,
        {"indonesia_freight_hkd": indonesia_freight_hkd},
    )
    return workbook, workbook.active


def _find_cell(sheet, value: object):
    return next(
        cell
        for row in sheet.iter_rows()
        for cell in row
        if cell.value == value
    )


@pytest.mark.parametrize(
    ("misc_ratio", "expected_settlement"),
    [("0.02", 0.98), ("0.03", 0.97), ("0.035", 0.965)],
)
def test_settlement_row_is_linked_to_the_exported_misc_ratio(
    misc_ratio: str,
    expected_settlement: float,
):
    workbook, sheet = _summary_sheet(
        rate_percent=None,
        material_amount_hkd="0",
        misc_ratio=misc_ratio,
    )
    try:
        assert sheet["Q6"].value == float(misc_ratio)
        assert 1 - sheet["Q6"].value == pytest.approx(expected_settlement)
        settlement_label = _find_cell(sheet, "÷")
        settlement_cell = sheet.cell(settlement_label.row, 4)
        assert settlement_cell.value == "=1-$Q$6"
        assert settlement_cell.number_format == "0.0000"
        quote_cell = sheet.cell(settlement_label.row + 1, 4)
        assert quote_cell.value.endswith(f"/D{settlement_label.row}")

        summary_header = _find_cell(sheet, "旺季价")
        summary_price_cell = sheet.cell(summary_header.row + 1, 4)
        assert summary_price_cell.value == f"=D{quote_cell.row}"

        cost_header_row = summary_header.row + 3
        misc_header = next(
            cell
            for cell in sheet[cost_header_row]
            if cell.value == "杂项"
        )
        misc_amount_cell = sheet.cell(misc_header.row + 1, misc_header.column)
        assert misc_amount_cell.value == f"=D{summary_price_cell.row}*$Q$6"
    finally:
        workbook.close()


@pytest.mark.parametrize(
    ("region_code", "effective_indonesia_freight_hkd", "expects_indonesia_freight"),
    [
        ("", "0", False),
        ("mainland", "0", False),
        ("indonesia", "1.25", True),
    ],
)
def test_direct_cost_rows_keep_additional_tax_and_effective_indonesia_freight_separate(
    region_code: str,
    effective_indonesia_freight_hkd: str,
    expects_indonesia_freight: bool,
):
    workbook, sheet = _summary_sheet(
        rate_percent=None,
        material_amount_hkd="0",
        region_code=region_code,
        indonesia_freight_hkd=effective_indonesia_freight_hkd,
        additional_tax_hkd="0.5",
    )
    try:
        direct_labels = [
            sheet.cell(row, 2).value
            for row in range(1, sheet.max_row + 1)
        ]
        assert direct_labels.count("附加税") == 1
        additional_tax_row = direct_labels.index("附加税") + 1
        assert sheet.cell(additional_tax_row, 4).value == 0.5

        assert (direct_labels.count("印尼运费") == 1) is expects_indonesia_freight
        if expects_indonesia_freight:
            indonesia_freight_row = direct_labels.index("印尼运费") + 1
            assert sheet.cell(indonesia_freight_row, 4).value == 1.25

        # Direct costs are part of the pre-markup subtotal exactly once.  They
        # are no longer written into a second detail row named "杂项".
        assert "杂项" not in direct_labels
        route_header = _find_cell(sheet, "运输方案")
        subtotal_row = route_header.row + 3
        subtotal_formula = sheet.cell(subtotal_row, 4).value
        assert subtotal_formula.startswith("=SUM(D")
        assert subtotal_formula.endswith(f":D{route_header.row})")

        settlement_label = _find_cell(sheet, "÷")
        assert sheet.cell(settlement_label.row, 4).value == "=1-$Q$6"

        summary_header = _find_cell(sheet, "旺季价")
        summary_price_cell = sheet.cell(summary_header.row + 1, 4)
        cost_header_row = summary_header.row + 3
        misc_header = next(
            cell
            for cell in sheet[cost_header_row]
            if cell.value == "杂项"
        )
        assert sheet.cell(misc_header.row + 1, misc_header.column).value == (
            f"=D{summary_price_cell.row}*$Q$6"
        )

        no_labor_header = next(
            cell
            for cell in sheet[cost_header_row]
            if cell.value == "不含人工成本"
        )
        no_labor_formula = sheet.cell(no_labor_header.row + 1, no_labor_header.column).value
        assert '"附加税"' in no_labor_formula
        assert '"印尼运费"' in no_labor_formula
    finally:
        workbook.close()


def test_non_rebate_factory_leaves_sewing_clothes_tax_rate_and_deduction_empty():
    workbook, sheet = _summary_sheet(rate_percent=None, material_amount_hkd="3.2500")
    try:
        top_label = _find_cell(sheet, "车缝物料退税")
        assert sheet.cell(top_label.row + 1, top_label.column).value == ""

        tax_label = next(
            cell
            for cell in sheet[_find_cell(sheet, "人民币外购件成本").row]
            if cell.value == "车缝物料退税"
        )
        tax_rate_row = tax_label.row + 1
        deduction_row = tax_label.row + 3
        assert sheet.cell(tax_rate_row, tax_label.column).value == ""
        assert sheet.cell(deduction_row, tax_label.column).value == ""

        material_detail = _find_cell(sheet, "车衣物料")
        labor_detail = _find_cell(sheet, "车衣人工")
        assert sheet.cell(material_detail.row, 1).value in (None, "")
        assert sheet.cell(material_detail.row, 4).value == 3.25
        assert sheet.cell(labor_detail.row, 1).value in (None, "")
        assert sheet.cell(labor_detail.row, 4).value == 9.25
    finally:
        workbook.close()


def test_rebate_factory_uses_rr2_material_amount_instead_of_total_sewing_cost():
    workbook, sheet = _summary_sheet(rate_percent="11.5000", material_amount_hkd="3.2500")
    try:
        top_label = _find_cell(sheet, "车缝物料退税")
        assert sheet.cell(top_label.row + 1, top_label.column).value == 0.115

        tax_label = next(
            cell
            for cell in sheet[_find_cell(sheet, "人民币外购件成本").row]
            if cell.value == "车缝物料退税"
        )
        tax_rate_row = tax_label.row + 1
        deduction_row = tax_label.row + 3
        assert sheet.cell(tax_rate_row, tax_label.column).value == 0.115
        assert sheet.cell(deduction_row, tax_label.column).value == (
            f"=3.25*{tax_label.column_letter}{tax_rate_row}"
        )

        material_detail = _find_cell(sheet, "车衣物料")
        labor_detail = _find_cell(sheet, "车衣人工")
        assert sheet.cell(material_detail.row, 1).value == "¥13%"
        assert sheet.cell(material_detail.row, 4).value == 3.25
        assert sheet.cell(labor_detail.row, 1).value in (None, "")
        assert sheet.cell(labor_detail.row, 4).value == 9.25
    finally:
        workbook.close()
