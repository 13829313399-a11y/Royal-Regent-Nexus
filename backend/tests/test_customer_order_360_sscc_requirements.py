"""360 carton-label evidence; all inline order and product details are synthetic."""
from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from app.services.huakang_a_order_legacy import schedule_parser as parser
from customer_order_360_private_samples import private_po_samples


def _parse(tmp_path: Path, instruction: str, *, suffix: str = ".pdf"):
    source = tmp_path / ("release" + suffix)
    source.write_bytes(b"parser text fixture")
    text = "\n".join([
        "PURCHASE ORDER RELEASE (MSL1) RL-765432-6 - REVISION 0",
        "Customer PO Number : 7654322",
        "Customer Release No. : RL76543226",
        "1 76543220069 Each 480",
        "Toy Fictional Sample Doll",
        "MS SSCC : 12345678901231",
        "Master Carton Qty : 4.0",
        "\f",
        "Planned Inspection Date : 8-Jan-2027 Factory Commit Date : 8-Jan-2027",
        "\f",
        "UPC Code: 123456789012",
        instruction,
    ])
    return parser._parse_po_text(source, text, 10, tmp_path)


def test_specific_sscc_carton_instruction_survives_generic_no_stickers(tmp_path):
    record = _parse(tmp_path, "\n".join([
        "No stickers -POE China -MS markings on 1 set of adjacent sides only -PO# (12345678901-2345), DCPI #",
        "123-45-6789, VCP/SSP: 4/1, SSCC Barcode labeled on master carton -Factory Name Example Carton Products",
        "(Sampletown) Ltd / ID 12345678",
    ]))
    assert record["outer_carton_sticker"] == "SSCC Barcode labeled on master carton"
    assert record["outer_carton_sticker_sources"] == [{
        "file_name": "release.pdf", "relative_path": "release.pdf", "page": 3,
        "raw_text": "SSCC Barcode labeled on master carton",
    }]
    assert record["barcode"] == "123456789012"
    assert "12345678901231" not in record["outer_carton_sticker"]
    assert "123456789012" not in record["outer_carton_sticker"]


@pytest.mark.parametrize("instruction", [
    "SSCC Barcode labeled on master carton",
    "SSCC barcode labelled on the outer carton",
    "SSCC labels required on master cartons",
    "SSCC label must be applied to the outer carton",
    "Apply SSCC label to master carton",
    "SSCC labels on the outer carton",
    "SSCC Barcode\nlabeled on master\ncarton",
])
def test_explicit_carton_label_requirement_keeps_its_exact_source(tmp_path, instruction):
    record = _parse(tmp_path, instruction)
    assert record["outer_carton_sticker"] == " ".join(instruction.split())
    assert record["outer_carton_sticker_sources"][0]["raw_text"] == instruction
    assert record["outer_carton_sticker_sources"][0]["page"] == 3


@pytest.mark.parametrize("instruction", [
    "No SSCC labels required on master cartons",
    "SSCC barcode must not be labeled on master carton",
    "Do not apply SSCC label to master carton",
    "Do not use SSCC labels on the outer carton",
    "SSCC labels on master carton are not required",
    "SSCC labels on master carton are optional",
    "Not required: SSCC labels on master carton",
    "Optional: SSCC labels on master carton",
    "No need to apply SSCC label to master carton",
    "No requirement for SSCC labels on master carton",
    "SSCC labels on master carton: NOT REQUIRED",
    "SSCC labels on master carton are not applicable",
    "No SSCC Barcode labeled on master carton",
    "Barcode on color box must be EAN or UPC code",
    "Barcode on outer carton must be ITF code",
    "SSCC label required on product / color box",
    "MS SSCC : 12345678901231",
    "UPC Code: 123456789012; EAN Code : 1234567890128",
    "SSCC code available. Label on master carton is UPC.",
    "",
])
def test_no_positive_carton_requirement_is_inferred(tmp_path, instruction):
    record = _parse(tmp_path, instruction)
    assert record["outer_carton_sticker"] == ""
    assert record["outer_carton_sticker_sources"] == []
    assert record["barcode"] == "123456789012"


def test_requirement_is_not_joined_across_pdf_pages(tmp_path):
    record = _parse(tmp_path, "SSCC Barcode labeled on\fmaster carton")
    assert record["outer_carton_sticker"] == ""


def test_converted_workbook_text_does_not_claim_a_pdf_page(tmp_path):
    record = _parse(tmp_path, "SSCC Barcode labeled on master carton", suffix=".xlsx")
    assert record["outer_carton_sticker_sources"][0]["page"] is None


def test_parsed_label_and_page_survive_preview_api_response_schema(tmp_path, monkeypatch):
    from app.schemas.customer_order import CustomerOrderLineOut
    from app.services import customer_order_unified as service

    record = _parse(tmp_path, "SSCC Barcode labeled on master carton")
    # The file worker uses indexed private names. Public lineage must restore
    # the uploaded name after parsing, including the nested source evidence.
    record["file_name"] = "001_release.pdf"
    record["relative_path"] = "001_release.pdf"
    record["outer_carton_sticker_sources"][0].update({
        "file_name": "001_release.pdf", "relative_path": "001_release.pdf",
    })
    monkeypatch.setattr(parser, "parse_po_file", lambda *_: record.copy())
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, headers in (
        (service.ITEM_SHEET, service.ITEM_HEADERS),
        (service.ORDER_SHEET, service.ORDER_HEADERS),
        (service.REVIEW_SHEET, service.ORDER_HEADERS),
    ):
        sheet = workbook.create_sheet(name)
        for column, value in enumerate(headers, 1):
            sheet.cell(3, column, value)
        sheet.cell(8, 1, "取消单")
    content = BytesIO()
    workbook.save(content)
    preview = service.create_unified_customer_preview(
        customer_code="360", factory_id="huakang-a", received_date="2026-10-19",
        po_files=[("release.pdf", b"test source")],
        schedule_file_name="schedule.xlsx", schedule_content=content.getvalue(),
    )
    # FastAPI's declared model must preserve this fact and its source rather
    # than silently dropping the optional parser field during serialization.
    row = CustomerOrderLineOut.model_validate(preview["rows"][0]).model_dump(mode="json")
    assert row["customer_label"] == "SSCC Barcode labeled on master carton"
    assert row["lineage"]["customer_label"] == "release.pdf · 第3页 · SSCC Barcode labeled on master carton"
    assert "001_" not in row["lineage"]["customer_label"]


def test_pdf_extraction_keeps_three_page_limit_and_real_page_numbers(tmp_path, monkeypatch):
    class Page:
        def __init__(self, number):
            self.number = number

        def extract_text(self):
            assert self.number <= 3, "the existing extraction budget must not expand"
            return f"physical page {self.number}"

    class Pdf:
        pages = [Page(number) for number in range(1, 6)]

        def __enter__(self):
            return self

        def __exit__(self, *_):
            return None

    monkeypatch.setattr(parser.pdfplumber, "open", lambda *_: Pdf())
    parser._extract_pdf.cache_clear()
    text, count = parser._extract_pdf(str(tmp_path / "bounded.pdf"), 1, 1)
    assert count == 5
    assert [page.strip() for page in text.split("\f")] == [
        "physical page 1", "physical page 2", "physical page 3",
    ]
    parser._extract_pdf.cache_clear()


@pytest.mark.parametrize("sample", private_po_samples())
def test_supplied_pdf_label_requirements_and_original_preservation(sample):
    source, expected = sample["source"], sample["expected"]
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    directory = sample["manifest"]["po_directory"]
    record = parser.parse_po_file(source, directory)
    assert record["outer_carton_sticker"] == expected["outer_carton_sticker"]
    assert record["outer_carton_sticker_sources"] == [
        dict(file_name=source.name, relative_path=str(source.relative_to(directory)),
             page=entry["page"], raw_text=entry["raw_text"])
        for entry in expected["outer_carton_sticker_sources"]
    ]
    assert record["barcode"] == expected["barcode"]
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
