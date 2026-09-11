from io import BytesIO
from decimal import Decimal
import pytest
from fastapi import HTTPException
from openpyxl import Workbook
from app.services.carton_opening_preview import parse_rows,options_from_json
from app.services.carton_opening_preview import inbound_time
from datetime import datetime

OPTIONS={"customer_name":"迪奇","warehouse":"A","snapshot_date":"2026-09-01","dimension_unit":"cm","currency":"CNY"}
HEADERS=["PO","货号","纸品／纸质","长","宽","高","仓位","期初库存数量","单价"]


def test_original_inbound_time_overrides_legacy_batch_date_and_preserves_time():
    content=workbook([["PO","I","A535B外箱",30,20,15,"A1",10,2,"B","2026-08-20 09:30"]],HEADERS+["仓库","原入库时间"])
    rows,_=parse_rows("stock.xlsx",content,options_from_json(OPTIONS))
    assert rows[0]["warehouse"]=="B"
    assert rows[0]["original_inbound_at"]=="2026-08-20T09:30:00+08:00"
    assert rows[0]["snapshot_date"]=="2026-08-20"


@pytest.mark.parametrize("raw,expected", [
    ("2026年8月1日", "2026-08-01T00:00:00+08:00"),
    ("2026/8/1 09:30:12", "2026-08-01T09:30:12+08:00"),
    (datetime(2026,8,1,9,30), "2026-08-01T09:30:00+08:00"),
    ("2026-07-31T23:30:00Z", "2026-08-01T07:30:00+08:00"),
    (46235.5, "2026-08-01T12:00:00+08:00"),
])
def test_row_time_formats_are_normalized_without_losing_clock_time(raw, expected):
    assert inbound_time(raw,0,"第3行")==expected


def test_1904_excel_epoch_keeps_time():
    assert inbound_time(44773.5,1,"第3行")=="2026-08-01T12:00:00+08:00"


def test_1904_workbook_with_unformatted_numeric_date_uses_workbook_epoch():
    from openpyxl.utils.datetime import MAC_EPOCH
    wb=Workbook();wb.epoch=MAC_EPOCH
    wb.active.append(HEADERS+["入库时间"])
    wb.active.append(["PO","I","A33外箱",30,20,15,"A1",10,2,44773.5])
    content=BytesIO();wb.save(content)
    rows,_=parse_rows("stock.xlsx",content.getvalue(),options_from_json({"dimension_unit":"cm"}))
    assert rows[0]["occurred_at"]=="2026-08-01T12:00:00+08:00"


@pytest.mark.parametrize("value", [None,"","8月1日","2026-02-30","2026-08-01 25:00","garbage 2026-08-01"])
def test_missing_or_invalid_row_date_reports_source(value):
    options={key:value for key,value in OPTIONS.items() if key!="snapshot_date"}
    content=workbook([["PO","I","A33外箱",30,20,15,"A1",10,2,value]],HEADERS+["入库时间"])
    with pytest.raises(HTTPException) as exc:
        parse_rows("stock.xlsx",content,options_from_json(options))
    assert "第 3 行" in exc.value.detail and "入库时间" in exc.value.detail


def test_row_time_with_native_excel_datetime_and_zero_row_without_date():
    options={key:value for key,value in OPTIONS.items() if key!="snapshot_date"}
    content=workbook([["PO","I","A33外箱",30,20,15,"A1",10,2,datetime(2026,8,1,9,30)],
                      ["PO","ZERO","A33外箱",30,20,15,"A1",0,2,None]],HEADERS+["入库时间"])
    rows,_=parse_rows("stock.xlsx",content,options_from_json(options))
    assert rows[0]["occurred_at"]=="2026-08-01T09:30:00+08:00"
    assert rows[1]["status"]=="ZERO"


def workbook(rows,headers=HEADERS):
    wb=Workbook();ws=wb.active;ws.title="期初库存导入"
    ws.append(["期初数据"]);ws.append(headers)
    for row in rows: ws.append(row)
    example=wb.create_sheet("填写示例");example.append(headers);example.append(["不会被导入"])
    stream=BytesIO();wb.save(stream);return stream.getvalue()


def test_original_columns_parse_box_card_and_zero_without_guessing_price():
    content=workbook([["PO-1","000123","A33+B外箱",30,20,15,"A1",100,2.5],
        ["PO-1","000123","B3B内箱",15,10,5,"A1",200,None],
        ["PO-1","000123","H5A平卡",30,20,0,"A1",50,.3],
        ["PO-1","000123","A33+B外箱",30,20,15,"A1",0,None]])
    rows,warnings=parse_rows("stock.xlsx",content,options_from_json(OPTIONS))
    assert len(rows)==4
    assert rows[0]["paper_quality"]=="A33+B" and rows[0]["item_no"]=="000123"
    assert rows[0]["amount"]==Decimal("250.00")
    assert rows[1]["unit_price"] is None and rows[1]["amount"] is None
    assert rows[2]["specification"]=="30 × 20 × 0 cm" and rows[2]["unit"]=="张"
    assert rows[3]["status"]=="ZERO"


@pytest.mark.parametrize("quantity",[None,-1,"NaN","Infinity","wrong"])
def test_invalid_opening_quantity_is_not_sourced_from_inbound_columns(quantity):
    with pytest.raises(HTTPException):
        parse_rows("stock.xlsx",workbook([["PO","ITEM","A33外箱",30,20,15,"A1",quantity,2,999]],HEADERS+["入库箱"]),options_from_json(OPTIONS))


def test_original_amount_is_ignored_but_opening_amount_is_checked():
    rows,_=parse_rows("stock.xlsx",workbook([["PO","ITEM","A33外箱",30,20,15,"A1",10,2,999,"2026年8月1日"]],HEADERS+["入库金额","日期"]),options_from_json(OPTIONS))
    assert rows[0]["amount"]==20 and rows[0]["snapshot_date"]=="2026-08-01"
    assert any("原入库" in warning for warning in rows[0]["warnings"])
    with pytest.raises(HTTPException) as exc:
        parse_rows("stock.xlsx",workbook([["PO","ITEM","A33外箱",30,20,15,"A1",10,2,999]],HEADERS+["原结余金额"]),options_from_json(OPTIONS))
    assert "不一致" in exc.value.detail


def test_eight_column_original_quality_header_is_supported_and_units_required():
    headers=HEADERS[:8];headers[2]="纸质"
    rows,_=parse_rows("stock.xlsx",workbook([["PO","ITEM","A33+B外箱",30,20,15,"A1",10]],headers),options_from_json(OPTIONS))
    assert rows[0]["unit_price"] is None and rows[0]["paper_quality"]=="A33+B"
    with pytest.raises(HTTPException): options_from_json({key:value for key,value in OPTIONS.items() if key!="dimension_unit"})
    with pytest.raises(HTTPException): options_from_json({**OPTIONS,"snapshot_date":"9月1日"})


@pytest.mark.parametrize("field,value",[("尺寸单位","in"),("币种","USD")])
def test_batch_settings_cannot_silently_override_explicit_row_fields(field,value):
    with pytest.raises(HTTPException) as exc:
        parse_rows("stock.xlsx",workbook([["PO","I","A33外箱",30,20,15,"A1",10,2,value]],HEADERS+[field]),options_from_json(OPTIONS))
    assert "不一致" in exc.value.detail


def test_published_template_empty_inputs_ignore_examples_and_example_rows_parse():
    from pathlib import Path
    from openpyxl import load_workbook
    template=Path(__file__).resolve().parents[2]/"public/templates/carton-history-inventory-import-template.xlsx"
    with pytest.raises(HTTPException) as exc:
        parse_rows(template.name,template.read_bytes(),options_from_json(OPTIONS))
    assert "未找到可导入数据" in exc.value.detail
    wb=load_workbook(template)
    for target_row,source_row in enumerate(range(6,10),3):
        for col,cell in enumerate(wb["填写示例"][source_row],1):
            wb["期初库存导入"].cell(target_row,col,cell.value)
    stream=BytesIO();wb.save(stream)
    rows,_=parse_rows(template.name,stream.getvalue(),options_from_json(OPTIONS))
    assert len(rows)==4 and sum(row["status"]=="ZERO" for row in rows)==1
    assert sum(row["unit_price"] is None for row in rows if row["status"]!="ZERO")==1
    assert [(row["packaging_type"],row["opening_quantity"]) for row in rows[:3]]==[("外箱",100),("内箱",200),("平卡",50)]
