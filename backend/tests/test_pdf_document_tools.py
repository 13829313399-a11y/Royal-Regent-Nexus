import sys
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter


BACKEND_DIR = Path(__file__).resolve().parents[1]

if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))


class FakeTable:
    bbox = (20, 100, 500, 180)

    def extract(self):
        return [["物料号", "数量"], ["00125", "8"]]


class FakeWordPage:
    def extract_text(self, **_kwargs):
        return "采购订单\n物料号 数量\n00125 8\n交货日期：2026-08-20"

    def find_tables(self, *, table_settings):
        assert table_settings["vertical_strategy"] == "lines"
        return [FakeTable()]

    def extract_words(self, **_kwargs):
        return [
            {"text": "采购订单", "x0": 20, "x1": 90, "top": 20, "bottom": 32},
            {"text": "物料号", "x0": 25, "x1": 80, "top": 115, "bottom": 128},
            {"text": "00125", "x0": 25, "x1": 70, "top": 145, "bottom": 158},
            {"text": "交货日期：", "x0": 20, "x1": 90, "top": 220, "bottom": 232},
            {"text": "2026-08-20", "x0": 95, "x1": 170, "top": 220, "bottom": 232},
        ]


class FakeWordDocument:
    pages = [FakeWordPage()]

    def close(self):
        return None


def _sample_pdf(page_count: int) -> bytes:
    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_blank_page(width=595, height=842)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_to_word_creates_editable_paragraphs_and_tables(monkeypatch):
    from app.services import pdf_to_word as service

    monkeypatch.setattr(service.pdfplumber, "open", lambda _stream: FakeWordDocument())

    result = service.convert_pdf_to_word(b"%PDF-fake", "../采购/PO-001.pdf")

    assert result.output_file_name == "PO-001_转换结果.docx"
    assert result.page_count == 1
    assert result.table_count == 1
    assert result.text_page_count == 1
    assert result.ocr_page_count == 0

    document = Document(BytesIO(result.content))
    assert [paragraph.text for paragraph in document.paragraphs if paragraph.text] == [
        "采购订单",
        "交货日期： 2026-08-20",
    ]
    assert document.tables[0].cell(1, 0).text == "00125"
    assert document.tables[0].cell(1, 1).text == "8"


def test_pdf_to_word_uses_ocr_when_native_content_is_missing(monkeypatch):
    from app.services import pdf_to_word as service

    class ScannedPage:
        def extract_text(self, **_kwargs):
            return ""

        def find_tables(self, **_kwargs):
            return []

        def extract_words(self, **_kwargs):
            return []

    class ScannedDocument(FakeWordDocument):
        pages = [ScannedPage()]

    monkeypatch.setattr(service.pdfplumber, "open", lambda _stream: ScannedDocument())
    monkeypatch.setattr(service.pdf_extraction, "_ocr_page", lambda _bytes, _index: [["OCR", "文字"]])

    result = service.convert_pdf_to_word(b"%PDF-scan", "扫描件.pdf")

    assert result.ocr_page_count == 1
    document = Document(BytesIO(result.content))
    assert document.paragraphs[0].text == "OCR 文字"


def test_pdf_split_supports_each_page_and_named_ranges():
    from app.services.pdf_split import split_pdf

    each_page = split_pdf(_sample_pdf(3), "订单.pdf", mode="each_page")
    assert each_page.output_file_name == "订单_拆分结果.zip"
    assert each_page.page_count == 3
    assert each_page.file_count == 3
    with ZipFile(BytesIO(each_page.content)) as archive:
        assert archive.namelist() == ["订单_第1页.pdf", "订单_第2页.pdf", "订单_第3页.pdf"]
        assert len(PdfReader(BytesIO(archive.read("订单_第2页.pdf"))).pages) == 1

    ranged = split_pdf(_sample_pdf(5), "计划.pdf", mode="ranges", page_ranges="1-2, 4, 5")
    assert ranged.file_count == 3
    with ZipFile(BytesIO(ranged.content)) as archive:
        assert archive.namelist() == ["计划_第1-2页.pdf", "计划_第4页.pdf", "计划_第5页.pdf"]
        assert len(PdfReader(BytesIO(archive.read("计划_第1-2页.pdf"))).pages) == 2


@pytest.mark.parametrize(
    ("page_ranges", "message"),
    [
        ("", "请输入"),
        ("3-1", "不是有效"),
        ("1-4", "超出"),
        ("第1页", "格式不正确"),
    ],
)
def test_pdf_split_rejects_invalid_ranges(page_ranges, message):
    from app.services.pdf_split import PdfSplitError, split_pdf

    with pytest.raises(PdfSplitError, match=message):
        split_pdf(_sample_pdf(3), "订单.pdf", mode="ranges", page_ranges=page_ranges)
