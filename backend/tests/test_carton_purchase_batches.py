"""Batch-confirmed procurement stays traceable without combining order quantities."""
from io import BytesIO
import json

from openpyxl import load_workbook

from test_molding_sample_api import make_client, login_as
from test_carton_procurement_api import _create_order, _order_payload
from test_carton_supplier_portal import BASE, setup_portal, supplier_login

PROCUREMENT = "/api/carton-procurement"


def submit_batch(client, orders, **changes):
    return client.post(PROCUREMENT + "/orders/bulk-submit-supplier", json={
        "factory_id": "huaxing", "items": [
            {"order_no": order["order_no"], "expected_revision": order["revision"]}
            for order in orders
        ], **changes,
    })


def supplier_documents(client):
    response = client.get(BASE + "/documents", params={"factory_id": "huaxing"})
    assert response.status_code == 200, response.text
    return response.json()


def workbook_rows(response):
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    rows = list(book.active.values)
    names = book.sheetnames
    book.close()
    return names, rows


def test_confirmed_batch_is_one_supplier_document_with_atomic_initial_snapshots(monkeypatch):
    with make_client(monkeypatch) as client:
        existing = setup_portal(client)
        login_as(client, "admin")
        drafts = [_create_order(client, submit_supplier=False) for _ in range(2)]
        response = submit_batch(client, [existing, *drafts])
        assert response.status_code == 200, response.text
        saved = response.json()
        assert len(saved) == 2
        batch = saved[0]["purchase_order_batch"]
        assert batch == saved[1]["purchase_order_batch"]
        assert batch["order_count"] == 2
        assert batch["document_no"].startswith("CG-")
        contexts = [client.get(PROCUREMENT + f"/orders/{order['order_no']}/purchase-order-context",
                               params={"factory_id": "huaxing"}).json() for order in saved]
        assert all(context["issues"][0]["purchase_order_batch"] == batch for context in contexts)
        assert all(context["pending_type"] == "NONE" for context in contexts)

        names, original_rows = workbook_rows(client.get(
            PROCUREMENT + f"/purchase-order-batches/{batch['id']}.xlsx",
            params={"factory_id": "huaxing"},
        ))
        assert names == ["合并采购单"]
        assert batch["document_no"] in original_rows[1][0]
        assert {row[3] for row in original_rows[5:9]} == {order["order_no"] for order in saved}
        assert all(row[13] > 0 for row in original_rows[5:9])
        assert client.get(PROCUREMENT + f"/purchase-order-batches/{batch['id']}.xlsx",
                          params={"factory_id": "huadeng"}).status_code == 404

        supplier_login(client)
        documents = supplier_documents(client)
        merged = next(row for row in documents if row["id"] == batch["id"])
        assert len(documents) == 2  # One earlier single issue plus this confirmed batch.
        assert merged["is_batch"] is True
        assert len(merged["orders"]) == 2 and len(merged["lines"]) == 4
        assert {line["order_no"] for line in merged["lines"]} == {order["order_no"] for order in saved}
        assert {line["source_document_no"] for line in merged["lines"]} == {
            context["issues"][0]["document_no"] for context in contexts}
        assert {line["unit"] for line in merged["lines"]} == {"个", "张"}
        workspace = client.get(BASE + "/workspace", params={"factory_id": "huaxing"}).json()
        batch_orders = [order for order in workspace["orders"] if order["id"] in {row["id"] for row in saved}]
        assert len(batch_orders) == 2
        assert all(not line["accepted"] for order in batch_orders for line in order["lines"])
        assert client.get(PROCUREMENT + f"/purchase-order-batches/{batch['id']}.xlsx",
                          params={"factory_id": "huaxing"}).status_code == 403

        selections = [{"factory_id": "huaxing", "kind": "PURCHASE", "id": batch["id"]}]
        names, rows = workbook_rows(client.post(BASE + "/documents/export.xlsx", json={"documents": selections}))
        assert names == ["采购单"] and len(rows) == 5
        assert all(row[1] == batch["document_no"] for row in rows[1:])
        assert {row[4] for row in rows[1:]} == {order["order_no"] for order in saved}
        assert {row[18] for row in rows[1:]} == {source["document_no"] for source in merged["source_documents"]}
        _, imported = workbook_rows(client.post(BASE + "/documents/order-import.xlsx", json={"documents": selections}))
        assert len(imported) == 5
        assert all(row[19] == batch["document_no"] for row in imported[1:])
        assert sum(row[6] for row in imported[1:]) == sum(float(line["change_quantity"]) for line in merged["lines"])
        assert next(row for row in supplier_documents(client) if row["id"] == batch["id"])["export_count"] == 2
        assert client.post(BASE + "/documents/export.xlsx", json={"documents": [
            {**selections[0], "factory_id": "huadeng"}]}).status_code in (403, 422)
        assert client.post(BASE + "/documents/export.xlsx", json={"documents": [
            {**selections[0], "id": contexts[0]["issues"][0]["id"]}]}).status_code == 404
        assert next(row for row in supplier_documents(client) if row["id"] == batch["id"])["export_count"] == 2

        # Acceptance still uses each exact paper and immutable original issue.
        accepted = client.put(BASE + "/commitments/batch", json={"factory_id": "huaxing", "lines": [
            {"order_line_id": line["id"], "issue_id": order["issue_id"],
             "expected_revision": line["commitment_revision"], "promised_date": "2026-10-10"}
            for order in batch_orders for line in order["lines"]
        ]})
        assert accepted.status_code == 200, accepted.text
        assert set(accepted.json()["order_ids"]) == {order["id"] for order in saved}

        login_as(client, "admin")
        first = saved[0]
        appended = client.post(PROCUREMENT + f"/orders/{first['order_no']}/append", json={
            "factory_id": "huaxing", "expected_revision": first["revision"], "additional_quantity": 120,
        })
        assert appended.status_code == 200, appended.text
        change = client.post(PROCUREMENT + f"/orders/{first['order_no']}/purchase-order-issues.xlsx", json={
            "factory_id": "huaxing", "expected_revision": appended.json()["revision"],
        })
        assert change.status_code == 200, change.text
        _, unchanged_rows = workbook_rows(client.get(
            PROCUREMENT + f"/purchase-order-batches/{batch['id']}.xlsx", params={"factory_id": "huaxing"}))
        assert unchanged_rows == original_rows
        supplier_login(client)
        documents = supplier_documents(client)
        assert next(row for row in documents if row["id"] == batch["id"])["lines"] == merged["lines"]
        assert any(row["document_type"] == "APPEND" and not row.get("is_batch") for row in documents)


def test_stale_batch_rolls_back_and_retry_does_not_create_another_batch(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        drafts = [_create_order(client, submit_supplier=False) for _ in range(2)]
        stale = [drafts[0], {**drafts[1], "revision": drafts[1]["revision"] + 1}]
        assert submit_batch(client, stale).status_code == 409
        listed = client.get(PROCUREMENT + "/orders", params={"factory_id": "huaxing", "limit": 100})
        assert listed.status_code == 200, listed.text
        states = {order["order_no"]: order for order in listed.json()["items"]}
        for order in drafts:
            assert states[order["order_no"]]["status"] == "CONFIRMED"
            context = client.get(PROCUREMENT + f"/orders/{order['order_no']}/purchase-order-context",
                                 params={"factory_id": "huaxing"}).json()
            assert context["issues"] == []
        saved = submit_batch(client, drafts)
        assert saved.status_code == 200, saved.text
        batch = saved.json()[0]["purchase_order_batch"]
        repeated = submit_batch(client, drafts)
        assert repeated.status_code == 200 and repeated.json() == []
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonAuditEvent, CartonPurchaseOrderIssue
        from sqlalchemy import select
        with SessionLocal() as db:
            issues = list(db.scalars(select(CartonPurchaseOrderIssue)).all())
            assert len(issues) == 2
            assert all(json.loads(issue.snapshot_json)["purchase_batch"]["id"] == batch["id"] for issue in issues)
            events = list(db.scalars(select(CartonAuditEvent).where(
                CartonAuditEvent.event_type == "PURCHASE_ORDER_BATCH_ISSUED")).all())
            assert len(events) == 1


def test_separate_confirmations_at_the_same_time_do_not_merge(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client)
        login_as(client, "admin")
        from app.services import carton_procurement as service
        monkeypatch.setattr(service, "now_text", lambda: "2026-09-29T12:00:00+08:00")
        first = [_create_order(client, submit_supplier=False) for _ in range(2)]
        second = [_create_order(client, submit_supplier=False) for _ in range(2)]
        one = submit_batch(client, first).json()[0]["purchase_order_batch"]
        two = submit_batch(client, second).json()[0]["purchase_order_batch"]
        assert one["generated_at"] == two["generated_at"]
        assert one["id"] != two["id"] and one["document_no"] != two["document_no"]
        supplier_login(client)
        batches = [row for row in supplier_documents(client) if row.get("is_batch")]
        assert {row["id"] for row in batches} == {one["id"], two["id"]}
        assert all(len(row["orders"]) == 2 for row in batches)


def test_confirmation_does_not_merge_different_suppliers(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        first = _create_order(client, submit_supplier=False)
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonSupplier
        with SessionLocal() as db:
            db.add(CartonSupplier(id="OTHER-SUPPLIER", factory_id="huaxing", supplier_code="OTHER",
                                 supplier_name="另一供应商", status="ACTIVE", created_at="2026-09-29", updated_at="2026-09-29"))
            db.commit()
        other = client.post(PROCUREMENT + "/orders", json={**_order_payload(), "supplier_id": "OTHER-SUPPLIER"})
        assert other.status_code == 201, other.text
        response = submit_batch(client, [first, other.json()])
        assert response.status_code == 200, response.text
        assert all(order["purchase_order_batch"] is None for order in response.json())


def test_issued_batch_blocks_member_deletion_and_survives_cancelled_tombstone(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client)
        login_as(client, "admin")
        response = submit_batch(client, [_create_order(client, submit_supplier=False) for _ in range(2)])
        assert response.status_code == 200, response.text
        saved = response.json()
        batch = saved[0]["purchase_order_batch"]
        path = PROCUREMENT + f"/purchase-order-batches/{batch['id']}.xlsx"
        params = {"factory_id": "huaxing"}
        _, original_rows = workbook_rows(client.get(path, params=params))
        for order in saved:
            assert not order["can_delete"]
            rejected = client.post(PROCUREMENT + f"/orders/{order['order_no']}/delete", json={
                **params, "expected_revision": order["revision"], "reason": "合并订单误删测试",
            })
            assert rejected.status_code == 409, rejected.text
            assert "合并采购单" in rejected.json()["detail"]
            assert workbook_rows(client.get(path, params=params))[1] == original_rows
        rejected = client.post(PROCUREMENT + "/orders/bulk-delete", json={
            **params, "reason": "合并订单批量误删测试", "items": [
                {"order_no": order["order_no"], "expected_revision": order["revision"]} for order in saved
            ],
        })
        assert rejected.status_code == 409, rejected.text
        anchor_order = saved[0]
        cancelled = client.post(PROCUREMENT + f"/orders/{anchor_order['order_no']}/reduce", json={
            **params, "expected_revision": anchor_order["revision"],
            "reduction_quantity": anchor_order["product_order_quantity"], "reason": "客户退回本张订单",
        })
        assert cancelled.status_code == 200, cancelled.text
        assert cancelled.json()["status"] == "CANCELLED"
        removed = client.post(PROCUREMENT + f"/orders/{anchor_order['order_no']}/delete", json={
            **params, "expected_revision": cancelled.json()["revision"], "reason": "取消后从台账移除",
        })
        assert removed.status_code == 204, removed.text
        assert workbook_rows(client.get(path, params=params))[1] == original_rows
        supplier_login(client)
        documents = supplier_documents(client)
        retained = next(row for row in documents if row["id"] == batch["id"])
        assert retained["is_batch"] and len(retained["orders"]) == 2
        assert {row["order_no"] for row in retained["orders"]} == {order["order_no"] for order in saved}
