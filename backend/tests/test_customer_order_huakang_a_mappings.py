from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from app.api import customer_order as customer_order_api
from app.services import customer_order_huakang_a as service
from app.services.huakang_a_order_legacy import green_toys_headstart


def _schedule_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "360客排期表"
    worksheet.cell(1, 1, "华康A 360客排期")
    headers = {
        4: "入单日期",
        5: "正单合同号",
        6: "产品货号",
        8: "产品名称",
        10: "PO数量",
        12: "验货日期",
        13: "走货期 FCD",
        20: "第三方客户 PO NO#",
        21: "跟单",
        22: "Customer Release No.",
        23: "外箱装箱数",
        24: "总箱数",
        35: "MS Container Type (FCL/LCL)",
        37: "Port of Discharge",
        40: "走货方式",
    }
    for column, value in headers.items():
        worksheet.cell(3, column, value)
    worksheet.merge_cells("A5:C5")
    worksheet.merge_cells("D5:H5")
    worksheet.cell(5, 1, "OLD-ITEM")
    worksheet.cell(5, 4, "旧产品")
    worksheet.cell(5, 1).font = openpyxl.styles.Font(
        name="微软雅黑", size=11, bold=True
    )
    worksheet.cell(5, 4).font = openpyxl.styles.Font(
        name="微软雅黑", size=11, bold=True
    )
    worksheet.cell(5, 1).alignment = openpyxl.styles.Alignment(
        horizontal="center", vertical="center"
    )
    worksheet.cell(5, 4).alignment = openpyxl.styles.Alignment(
        horizontal="left", vertical="center"
    )
    worksheet.row_dimensions[5].height = 28
    worksheet.cell(6, 5, "OLD-HISTORY")
    worksheet.cell(6, 6, "OLD-ITEM")
    worksheet.cell(6, 8, "旧产品")
    worksheet.cell(6, 10, 1)
    for column in range(1, 41):
        worksheet.cell(6, column).font = openpyxl.styles.Font(
            name="微软雅黑", size=9
        )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _parsed_record(file_name: str = "RL-100-1.pdf") -> dict[str, object]:
    return {
        "file_name": file_name,
        "contract_no": "RL-100-1",
        "release_type": "FIRM",
        "revision": 2,
        "revision_date": "03-Jul-2026",
        "customer_po": "PO-360-1",
        "customer_release": "CR-1",
        "contact": "Amy",
        "container_type": "FCL",
        "item_no": "60350",
        "description": "ThreeSixty Toy",
        "quantity": 240,
        "master_carton_qty": 12,
        "planned_inspection_date": "2026-07-10",
        "factory_commit_date": "2026-07-17",
        "transportation_mode": "SEA",
        "port": "FELIXSTOWE",
        "parse_ok": True,
        "parse_warnings": [],
    }


def _green_toys_schedule_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "2026年排期"
    headers = {
        1: "来单\n日期",
        2: "PO",
        3: "客户名称",
        4: "产品货号",
        5: "出货产品货号",
        6: "产品名称",
        7: "描述",
        8: "订单数量",
        10: "客验日期",
        12: "走货期",
        16: "走货\n国家",
        17: "落货港",
        20: "外箱\n装箱数",
        21: "总箱数",
        22: "走货方式",
        23: "单价US",
        24: "货价总金额US$",
        25: "PO单价HK$",
        26: "货价总金额HK$",
        29: "系统",
    }
    for column, value in headers.items():
        worksheet.cell(3, column, value)
    worksheet.merge_cells("A4:D4")
    worksheet.cell(4, 1, "BBMSUB-1839 潜艇故事书")
    worksheet.cell(5, 1, datetime(2026, 7, 1))
    worksheet.cell(5, 2, "7000-1")
    worksheet.cell(5, 3, "GT")
    worksheet.cell(5, 4, "BBMSUB-1839")
    worksheet.cell(5, 5, "BBMSUB-1839")
    worksheet.cell(5, 6, "潜艇故事书")
    worksheet.cell(5, 7, "Bubbling Submarine & Board Book")
    worksheet.cell(5, 8, 1000)
    worksheet.cell(5, 16, "美國")
    worksheet.cell(5, 17, "Hayward, CA 94545 USA")
    worksheet.cell(5, 20, 8)
    worksheet.cell(5, 21, "=H5/T5")
    worksheet.cell(5, 23, 1.449)
    for column in range(1, 38):
        worksheet.cell(5, column).font = openpyxl.styles.Font(name="宋体", size=9)
    audit = workbook.create_sheet("测试保留页")
    audit["A1"] = "keep"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _headstart_schedule_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "2026年HS客排期表 "
    headers = {
        1: "客户",
        2: "入单\n日期",
        3: "正单合同号",
        4: "产品货号",
        5: "产品名称",
        7: "PO数量",
        11: "验货日期",
        12: "走货期",
        17: "走货\n国家",
        18: "客户名称Keycode",
        20: "跟单",
        22: "外箱\n装箱数",
        24: "总箱数",
        25: "日期码",
        26: "条码",
        27: "订单单价USD",
        28: "总金额USD",
        29: "单价HK$",
        30: "总金额HK$",
        33: "系统",
        34: "出货货号",
    }
    for column, value in headers.items():
        worksheet.cell(6, column, value)
    worksheet.merge_cells("B7:E7")
    worksheet.cell(7, 1, "78560")
    worksheet.cell(7, 2, "Disney Giftables Series 1")
    worksheet.cell(8, 1, "Headstart")
    worksheet.cell(8, 3, "OLD-HS")
    worksheet.cell(8, 4, "78560")
    worksheet.cell(8, 5, "Disney Giftables Series 1 一代迪士尼公仔")
    worksheet.cell(8, 7, 6000)
    worksheet.cell(8, 17, "澳洲")
    worksheet.cell(8, 18, "HSINT AU")
    worksheet.cell(8, 20, "Billie")
    worksheet.cell(8, 22, 12)
    worksheet.cell(8, 24, "=G8/V8")
    worksheet.cell(8, 27, 2.14)
    for column in range(1, 46):
        worksheet.cell(8, column).font = openpyxl.styles.Font(name="宋体", size=9)
    audit = workbook.create_sheet("测试保留页")
    audit["A1"] = "keep"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_huakang_a_360_registry_coexists_with_huaxing_360() -> None:
    assert set(service.HUAKANG_A_CUSTOMER_MAPPINGS) == {
        "360",
        "green-toys",
        "headstart",
    }
    assert customer_order_api.CUSTOMER_FACTORY_IDS["360"] == "huaxing"
    assert customer_order_api.CUSTOMER_FACTORY_OPTIONS["360"] == (
        "huaxing",
        "huakang-a",
    )
    assert customer_order_api._get_mapped_customer_spec(
        "360", "huaxing"
    ).target_template == "HUAXING_360_SCHEDULE_APPEND_V2"
    assert customer_order_api._get_mapped_customer_spec(
        "360", "huakang-a"
    ).target_template == "HUAKANG_A_360_SCHEDULE_APPEND_V3"
    assert customer_order_api.CUSTOMER_FACTORY_OPTIONS["green-toys"] == (
        "huakang-a",
    )
    assert customer_order_api.CUSTOMER_FACTORY_OPTIONS["headstart"] == (
        "huakang-a",
    )

    with pytest.raises(service.HuakangACustomerOrderError, match="只属于华康A厂区"):
        service.create_huakang_a_customer_preview(
            customer_code="360",
            factory_id="huakang-c",
            received_date="2026-08-03",
            po_files=[("RL-100-1.pdf", b"pdf")],
            schedule_file_name="360排期.xlsx",
            schedule_content=_schedule_bytes(),
        )


def test_legacy_parser_extracts_threesixty_release_fields(tmp_path: Path) -> None:
    source = tmp_path / "RL-100-1.txt"
    source.write_text("placeholder", encoding="utf-8")
    text = """
PURCHASE ORDER RELEASE (FIRM) RL-100-1
REVISION 2
Revision Date : 03-Jul-2026
Customer PO Number : PO-360-1
Customer Release No. : CR-1
Our Contact : Amy
MS Container Type (FCL/LCL) : FCL
1 60350 Each 240
ThreeSixty Toy
Master Carton Qty : 12
Planned Inspection Date : 10-Jul-2026 Factory Commit Date : 17-Jul-2026
Transportation Mode : SEA
Port of Discharge : FELIXSTOWE
"""

    parsed = service.schedule_parser._parse_po_text(source, text, 1, tmp_path)

    assert parsed["parse_ok"] is True
    assert parsed["contract_no"] == "RL-100-1"
    assert parsed["revision"] == 2
    assert parsed["item_no"] == "60350"
    assert parsed["quantity"] == 240
    assert parsed["master_carton_qty"] == 12
    assert parsed["planned_inspection_date"] == "2026-07-10"
    assert parsed["factory_commit_date"] == "2026-07-17"


def test_huakang_a_360_preview_and_export_follow_legacy_mapping(monkeypatch) -> None:
    monkeypatch.setattr(
        service.schedule_parser,
        "parse_po_file",
        lambda path, _root: _parsed_record(Path(path).name),
    )
    schedule = _schedule_bytes()
    preview = service.create_huakang_a_customer_preview(
        customer_code="360",
        factory_id="huakang-a",
        received_date="2026-08-03",
        po_files=[("RL-100-1.pdf", b"pdf")],
        schedule_file_name="2026年华康A 360排期.xlsx",
        schedule_content=schedule,
    )

    row = preview["rows"][0]
    assert preview["factory_id"] == "huakang-a"
    assert preview["preview_schema_version"] == service.PREVIEW_SCHEMA_VERSION
    assert row["contract_no"] == "RL-100-1"
    assert row["po_no"] == "PO-360-1"
    assert row["product_no"] == "60350"
    assert row["product_name_en"] == "ThreeSixty Toy"
    assert row["carton_count"] == "20"
    assert row["line_q"] == "2026-07-10"
    assert row["requested_ship_date"] == "2026-07-17"
    assert row["status"] == "valid"

    output, file_name, exported_preview = service.export_huakang_a_customer_schedule(
        customer_code="360",
        factory_id="huakang-a",
        received_date="2026-08-03",
        po_files=[("RL-100-1.pdf", b"pdf")],
        schedule_file_name="2026年华康A 360排期.xlsx",
        schedule_content=schedule,
    )
    workbook = openpyxl.load_workbook(BytesIO(output), data_only=False)
    try:
        assert workbook.sheetnames == ["360客排期表"]
        worksheet = workbook["360客排期表"]
        assert worksheet.cell(6, 5).value == "OLD-HISTORY"
        assert worksheet.cell(7, 1).value == "60350"
        assert worksheet.cell(7, 4).value == "ThreeSixty Toy"
        assert {str(value) for value in worksheet.merged_cells.ranges} >= {
            "A7:C7",
            "D7:H7",
        }
        assert worksheet.row_dimensions[7].height == worksheet.row_dimensions[5].height
        assert worksheet.cell(7, 1).font.name == worksheet.cell(5, 1).font.name
        assert worksheet.cell(7, 1).font.sz == worksheet.cell(5, 1).font.sz
        assert worksheet.cell(7, 1).font.bold == worksheet.cell(5, 1).font.bold
        assert worksheet.cell(7, 4).font.name == worksheet.cell(5, 4).font.name
        assert worksheet.cell(7, 4).font.sz == worksheet.cell(5, 4).font.sz
        assert worksheet.cell(7, 4).font.bold == worksheet.cell(5, 4).font.bold
        assert worksheet.cell(8, 4).value == datetime(2026, 7, 3)
        assert worksheet.cell(8, 5).value == "RL-100-1"
        assert worksheet.cell(8, 6).value == "60350"
        assert worksheet.cell(8, 10).value == 240
        assert worksheet.cell(8, 24).value == 20
    finally:
        workbook.close()
    assert file_name == preview["output_file_name"]
    assert exported_preview["summary"] == {
        "total": 1,
        "valid": 1,
        "warning": 0,
        "blocked": 0,
    }


def test_missing_fcd_warns_but_missing_key_fields_block(monkeypatch) -> None:
    missing_fcd = _parsed_record()
    missing_fcd["factory_commit_date"] = ""
    missing_fcd["parse_warnings"] = ["未识别 FCD"]
    monkeypatch.setattr(
        service.schedule_parser,
        "parse_po_file",
        lambda _path, _root: missing_fcd,
    )
    preview = service.create_huakang_a_customer_preview(
        customer_code="360",
        factory_id="huakang-a",
        received_date="2026-08-03",
        po_files=[("RL-100-1.pdf", b"pdf")],
        schedule_file_name="360排期.xlsx",
        schedule_content=_schedule_bytes(),
    )
    assert preview["rows"][0]["status"] == "warning"

    missing_contract = dict(missing_fcd, contract_no="", parse_ok=False)
    missing_contract["parse_warnings"] = ["未识别合同号"]
    monkeypatch.setattr(
        service.schedule_parser,
        "parse_po_file",
        lambda _path, _root: missing_contract,
    )
    blocked = service.create_huakang_a_customer_preview(
        customer_code="360",
        factory_id="huakang-a",
        received_date="2026-08-03",
        po_files=[("bad.pdf", b"pdf")],
        schedule_file_name="360排期.xlsx",
        schedule_content=_schedule_bytes(),
    )
    assert blocked["rows"][0]["status"] == "blocked"


def test_headstart_text_parser_extracts_order_line_and_business_dates(monkeypatch) -> None:
    text = """
07−AUG−26
2022802AB
PURCHASE ORDER
Account Code Currency Order Terms
NP ROYALREGEN USD
Line Total Unit Cost Qty (Pieces) Item Description Item Code APN Number
24962−0336 BFF Single Pack CDU Wave 3 57996 0.737 42,743.05 21 DEC 2027 840150303756 36.00 1611
> Date Code : RR011226GC
> Legal line : year 2026
"""
    monkeypatch.setattr(
        green_toys_headstart,
        "_extract_headstart_text",
        lambda _content: text,
    )

    rows = green_toys_headstart.parse_headstart_pdf("headstart.pdf", b"pdf")

    assert len(rows) == 1
    assert rows[0]["contract_no"] == "2022802AB"
    assert rows[0]["item_no"] == "24962-0336"
    assert rows[0]["quantity"] == 57996
    assert rows[0]["planned_inspection_date"] == "2027-12-21"
    assert rows[0]["factory_commit_date"] == "2027-12-28"
    assert rows[0]["barcode"] == "840150303756"
    assert rows[0]["date_code"] == "RR011226GC"


def test_green_toys_ocr_parser_extracts_rows_and_normalizes_prices() -> None:
    text = """
Purchase Order
PO Date 7/17/2026
Deliver By Date 10/1/2026
Item Description Order Quantity Received Inventory Detail Unit Price Amount
BBIDTR-1850     John Deere Tractor & Board Book     5,000|     0     2.788     13,945.00
RHWC-1852       Rainbow Harvest - Color-Changing Watering Can Activity Set (Blue)     3,000|     0     441     13,27
"""

    rows = green_toys_headstart.parse_green_toys_ocr_text("7816.png", text)

    assert [row["customer_po"] for row in rows] == ["7816-1", "7816-2"]
    assert rows[0]["item_no"] == "BBIDTR-1850"
    assert rows[0]["quantity"] == 5000
    assert rows[0]["unit_price_usd"] == 2.789
    assert rows[0]["amount_usd"] == 13945.0
    assert rows[1]["unit_price_usd"] == 4.41
    assert rows[1]["amount_usd"] == 13230.0
    assert rows[0]["planned_inspection_date"] == "2026-09-30"
    assert rows[0]["factory_commit_date"] == "2026-10-01"


def test_green_toys_preview_inherits_master_data_and_exports_in_item_group(
    monkeypatch,
) -> None:
    text = """
Purchase Order
PO Date 7/20/2026
PO# 7818
Deliver By Date 9/22/2026
BBMSUB-1839     Bubbling Submarine & Board Book     2,500     0     1.449     3,622.50
"""
    monkeypatch.setattr(
        green_toys_headstart,
        "_ocr_green_toys_image",
        lambda _path: text,
    )
    schedule = _green_toys_schedule_bytes()

    preview = service.create_huakang_a_customer_preview(
        customer_code="green-toys",
        factory_id="huakang-a",
        received_date="2026-08-24",
        po_files=[("7818.png", b"png")],
        schedule_file_name="Green Toys排期.xlsx",
        schedule_content=schedule,
    )

    row = preview["rows"][0]
    assert row["status"] == "warning"
    assert row["product_name_zh"] == "潜艇故事书"
    assert row["units_per_carton"] == "8"
    assert row["country"] == "美國"
    assert row["requested_ship_date"] == "2026-09-22"

    output, _, _ = service.export_huakang_a_customer_schedule(
        customer_code="green-toys",
        factory_id="huakang-a",
        received_date="2026-08-24",
        po_files=[("7818.png", b"png")],
        schedule_file_name="Green Toys排期.xlsx",
        schedule_content=schedule,
    )
    workbook = openpyxl.load_workbook(BytesIO(output), data_only=False)
    try:
        assert workbook.sheetnames == ["2026年排期", "测试保留页"]
        worksheet = workbook["2026年排期"]
        assert worksheet.cell(6, 1).value == "BBMSUB-1839 潜艇故事书"
        assert worksheet.cell(7, 2).value == "7818-1"
        assert worksheet.cell(7, 4).value == "BBMSUB-1839"
        assert worksheet.cell(7, 8).value == 2500
        assert worksheet.cell(7, 20).value == 8
        assert worksheet.cell(7, 21).value == "=CEILING(H7/T7,1)"
        assert worksheet.cell(7, 24).value == "=W7*H7"
        assert worksheet.cell(7, 25).value == "=W7*7.8"
        assert worksheet.cell(7, 1).fill.fgColor.rgb == "00CCFFCC"
        assert workbook["测试保留页"]["A1"].value == "keep"
    finally:
        workbook.close()


def test_headstart_preview_and_export_create_new_item_title(monkeypatch) -> None:
    parsed = [{
        "file_name": "headstart.pdf",
        "contract_no": "2022999AA",
        "customer_po": "",
        "order_date": "2026-08-07",
        "item_no": "24962-0336",
        "description": "BFF Single Pack CDU Wave 3",
        "quantity": 3600,
        "master_carton_qty": 36,
        "total_cartons": 100,
        "planned_inspection_date": "2026-11-20",
        "factory_commit_date": "2026-11-27",
        "unit_price_usd": 0.737,
        "amount_usd": 2653.2,
        "barcode": "840150303756",
        "date_code": "RR011226GC",
        "account_code": "NP",
        "country": "澳洲",
        "packaging": "Legal line : year 2026",
        "parse_warnings": [],
        "parse_ok": True,
    }]
    monkeypatch.setattr(
        green_toys_headstart,
        "parse_headstart_pdf",
        lambda _file_name, _content: [dict(parsed[0])],
    )
    schedule = _headstart_schedule_bytes()

    preview = service.create_huakang_a_customer_preview(
        customer_code="headstart",
        factory_id="huakang-a",
        received_date="2026-08-24",
        po_files=[("headstart.pdf", b"pdf")],
        schedule_file_name="HeadStart排期.xlsx",
        schedule_content=schedule,
    )
    assert preview["summary"] == {
        "total": 1,
        "valid": 1,
        "warning": 0,
        "blocked": 0,
    }

    output, _, _ = service.export_huakang_a_customer_schedule(
        customer_code="headstart",
        factory_id="huakang-a",
        received_date="2026-08-24",
        po_files=[("headstart.pdf", b"pdf")],
        schedule_file_name="HeadStart排期.xlsx",
        schedule_content=schedule,
    )
    workbook = openpyxl.load_workbook(BytesIO(output), data_only=False)
    try:
        worksheet = workbook["2026年HS客排期表 "]
        assert worksheet.cell(9, 1).value == "24962-0336"
        assert worksheet.cell(9, 2).value == "BFF Single Pack CDU Wave 3"
        assert "B9:E9" in {str(value) for value in worksheet.merged_cells.ranges}
        assert worksheet.cell(10, 3).value == "2022999AA"
        assert worksheet.cell(10, 4).value == "24962-0336"
        assert worksheet.cell(10, 24).value == "=G10/V10"
        assert worksheet.cell(10, 28).value == "=AA10*G10"
        assert worksheet.cell(10, 29).value == "=AA10*7.8"
        assert worksheet.cell(10, 1).fill.fgColor.rgb == "00CCFFCC"
        assert workbook["测试保留页"]["A1"].value == "keep"
    finally:
        workbook.close()
