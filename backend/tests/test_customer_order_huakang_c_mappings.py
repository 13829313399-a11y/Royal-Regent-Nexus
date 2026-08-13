from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Border, Font, PatternFill, Side

from app.api import customer_order as customer_order_api
from app.services import customer_order_huakang_c as service


def _schedule_bytes(customer_code: str) -> bytes:
    spec = service.get_huakang_c_customer_mapping(customer_code)
    profile = service.huakang_schedule.PROFILES[spec.legacy_code]
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = profile.sheets[0]
    for row in range(1, profile.header_rows + 1):
        worksheet.cell(row, 1, f"{spec.name}表头{row}")
    item_columns, product_column = service.huakang_schedule._product_columns(profile)
    worksheet.cell(profile.style_row, profile.key_cols[0], "OLD-CONTRACT")
    worksheet.cell(profile.style_row, profile.key_cols[1], "OLD-ITEM")
    worksheet.cell(profile.style_row, item_columns[0], "ITEM-1")
    worksheet.cell(profile.style_row, product_column, "排期标准品名")
    thin = Side(style="thin", color="000000")
    if customer_code == "jazwares":
        for col_no in range(1, profile.max_col + 1):
            cell = worksheet.cell(profile.style_row, col_no)
            cell.font = Font(name="Calibri", size=12)
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        worksheet.cell(profile.style_row, 10).font = Font(name="宋体", size=12)
        worksheet.cell(profile.style_row, 11, 1)
        worksheet.cell(profile.style_row, 25, f"=K{profile.style_row}/12")
        worksheet.cell(profile.style_row, 42, "=2*7.75")
        worksheet.cell(profile.style_row, 43, f"=AP{profile.style_row}*K{profile.style_row}")
        for col_no in (1, 2, 41, 44):
            worksheet.cell(profile.style_row, col_no).number_format = "yyyy/m/d;@"
        for col_no in (42, 43):
            worksheet.cell(profile.style_row, col_no).number_format = 'HK$#,##0.00'
        worksheet.row_dimensions[profile.style_row].height = 32

        summary_row = profile.style_row + 1
        worksheet.cell(summary_row, 10, "ITEM-1 合计")
        worksheet.cell(summary_row, 11, f"=SUM(K{profile.style_row}:K{profile.style_row})")
        for col_no in (10, 11):
            cell = worksheet.cell(summary_row, col_no)
            cell.font = Font(name="Calibri", size=12, bold=True, color="FF0000")
            cell.border = Border(left=thin, right=thin, top=thin)
        worksheet.row_dimensions[summary_row].height = 32
    elif customer_code == "maxx":
        headers = {
            1: "接单日期", 2: "放产日期", 3: "放产单号", 4: "客PO号",
            5: "合同", 6: "客名", 7: "产品品牌", 8: "大货号", 9: "小货号",
            10: "货名", 11: "订单数量", 12: "箱数", 25: "装箱", 28: "备注/国家",
            42: "订单截数期", 43: "单价HK$", 44: "金额HK$", 46: "走货方式",
            51: "备注", 57: "备注",
        }
        for col_no, value in headers.items():
            worksheet.cell(1, col_no, value)
        for col_no in range(1, profile.max_col + 1):
            cell = worksheet.cell(profile.style_row, col_no)
            cell.font = Font(name="Calibri", size=10)
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        worksheet.cell(profile.style_row, 11, "=120-120")
        worksheet.cell(profile.style_row, 12, f"=K{profile.style_row}/4")
        worksheet.cell(profile.style_row, 43, "=2*7.75")
        worksheet.cell(profile.style_row, 44, f"=+AQ{profile.style_row}*K{profile.style_row}")
        worksheet.row_dimensions[profile.style_row].height = 25
        second_detail_row = profile.style_row + 2
        for col_no in range(1, profile.max_col + 1):
            worksheet.cell(second_detail_row, col_no)._style = deepcopy(
                worksheet.cell(profile.style_row, col_no)._style
            )
        worksheet.cell(second_detail_row, 5, "OLD-CONTRACT-2")
        worksheet.cell(second_detail_row, 8, "ITEM-2")
        worksheet.cell(second_detail_row, 10, "排期另一个品名")
        worksheet.cell(second_detail_row, 11, "=200-200")
        worksheet.cell(second_detail_row, 12, f"=K{second_detail_row}/4")
        worksheet.cell(second_detail_row, 43, "=3*7.75")
        worksheet.cell(second_detail_row, 44, f"=+AQ{second_detail_row}*K{second_detail_row}")
        worksheet.row_dimensions[second_detail_row].height = 25
        for summary_row in (3, 5, 6, 7, 8):
            worksheet.cell(summary_row, 10, "合计")
            worksheet.cell(summary_row, 11, "=SUM(#REF!)")
            for col_no in range(1, profile.max_col + 1):
                cell = worksheet.cell(summary_row, col_no)
                cell.font = Font(name="Calibri", size=10, bold=col_no in (10, 11))
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            for col_no in (10, 11):
                worksheet.cell(summary_row, col_no).fill = PatternFill("solid", fgColor="F8CBAD")
            worksheet.row_dimensions[summary_row].height = 25
    elif customer_code == "index":
        headers = {
            1: "接单日期", 4: "客PO号", 5: "合同", 6: "客名", 7: "国家",
            8: "大货号", 9: "小货号", 10: "货名", 11: "订单数量",
            24: "装箱", 25: "纸箱", 41: "订单截数期", 42: "单价HK$",
            43: "金额HK$", 44: "出货期", 49: "备注",
        }
        for col_no, value in headers.items():
            worksheet.cell(1, col_no, value)
        for col_no in range(1, profile.max_col + 1):
            cell = worksheet.cell(profile.style_row, col_no)
            cell.font = Font(name="Calibri", size=10)
            cell.fill = PatternFill(fill_type=None)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        worksheet.cell(profile.style_row, 25, f"=K{profile.style_row}/X{profile.style_row}")
        worksheet.cell(profile.style_row, 42, "=2*7.75")
        worksheet.cell(profile.style_row, 43, f"=AP{profile.style_row}*K{profile.style_row}")
        worksheet.row_dimensions[profile.style_row].height = 36
    elif customer_code == "strottman":
        headers = {
            1: "能否备料", 2: "能否生产", 3: "款号", 4: "入单日期",
            5: "正单合同号", 6: "产品货号", 7: "出货货号", 8: "产品名称",
            9: "版本", 10: "PO数量", 11: "订单状态", 12: "验货日期",
            13: "走货期 FCD", 18: "走货国家", 19: "客户名称Keycode",
            21: "跟单", 22: "Customer Release No.", 23: "外箱装箱数",
            24: "总箱数", 25: "日期码", 26: "条码", 27: "订单单价USD",
            28: "单价HK$", 29: "总金额USD", 30: "总金额HK$",
            31: "出厂价HK$", 32: "出厂价总金额HK$", 33: "备注", 34: "系统",
        }
        for col_no, value in headers.items():
            worksheet.cell(3, col_no, value)
        worksheet["A9"] = "ITEM-1"
        worksheet["D9"] = "排期标准品名"
        for col_no in range(1, profile.max_col + 1):
            title_cell = worksheet.cell(9, col_no)
            title_cell.font = Font(name="宋体", size=9, bold=col_no in (1, 4))
            title_cell.border = Border(bottom=thin)
            detail_cell = worksheet.cell(10, col_no)
            detail_cell.font = Font(name="宋体", size=9)
            detail_cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        worksheet.merge_cells("A9:C9")
        worksheet.merge_cells("D9:K9")
        worksheet["A10"] = "正常备料"
        worksheet["B10"] = "正常生产"
        worksheet["C10"] = "ITEM-1"
        worksheet["D10"] = "2026-07-01"
        worksheet["D10"].number_format = "yyyy/m/d;@"
        worksheet["E10"] = "OLD-CONTRACT"
        worksheet["F10"] = "OLD-ITEM"
        worksheet["G10"] = "ITEM-1"
        worksheet["H10"] = "排期标准品名"
        worksheet["J10"] = 120
        worksheet["W10"] = 12
        worksheet["X10"] = "=J10/W10"
        worksheet["AA10"] = 2
        worksheet["AB10"] = "=AA10*7.75"
        worksheet["AC10"] = "=AA10*J10"
        worksheet["AD10"] = "=AB10*J10"
        worksheet["AF10"] = "=AE10*J10"
        worksheet["AH10"].number_format = "yyyy/m/d;@"
        worksheet.row_dimensions[9].height = 21
        worksheet.row_dimensions[10].height = 32
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _order(customer_code: str) -> dict[str, object]:
    spec = service.get_huakang_c_customer_mapping(customer_code)
    return {
        "customer_code": spec.legacy_code,
        "customer_name": spec.name,
        "po_number": "PO-100",
        "contract_no": "SC-100",
        "po_date": "2026-07-01",
        "ship_date": "2026-09-01",
        "ship_dates": ["2026-09-01"],
        "ship_to": "Hong Kong",
        "project_no": "",
        "project_name": "",
        "delivery_term": "",
        "version": "2" if customer_code == "jazwares" else "",
        "contact": "Amy",
        "lines": [{
            "item_code": "ITEM-1",
            "description": "PO Product",
            "qty": 120,
            "unit": "PCS",
            "unit_price_usd": 2.0,
            "amount_usd": 240.0,
            "pcs_per_carton": 12,
            "carton_qty": 10,
            "source_unit_price_usd": 0,
            "unit_price_hkd": 0,
            "version": "2" if customer_code == "jazwares" else "",
        }],
        "warnings": [],
    }


def test_huakang_c_registry_is_factory_scoped_and_keeps_maxx_separate() -> None:
    assert set(service.HUAKANG_C_CUSTOMER_MAPPINGS) == {
        "index", "jazwares", "maxx", "strottman", "jp",
    }
    assert customer_order_api.CUSTOMER_FACTORY_IDS["index"] == "huakang-c"
    assert customer_order_api.CUSTOMER_FACTORY_IDS["maxx"] == "huaxing"
    assert customer_order_api.CUSTOMER_FACTORY_OPTIONS["maxx"] == (
        "huaxing",
        "huakang-c",
    )
    assert customer_order_api._get_mapped_customer_spec(
        "maxx", "huaxing"
    ).target_template == "HUAXING_MAXX_SCHEDULE_APPEND_V2"
    assert customer_order_api._get_mapped_customer_spec(
        "maxx", "huakang-c"
    ).target_template == "HUAKANG_C_MAXX_SCHEDULE_APPEND_V2"

    with pytest.raises(service.HuakangCCustomerOrderError, match="只属于华康C厂区"):
        service.create_huakang_c_customer_preview(
            customer_code="index",
            factory_id="huakang-a",
            received_date="2026-08-03",
            po_files=[("INDEX.pdf", b"pdf")],
            schedule_file_name="INDEX排期.xlsx",
            schedule_content=_schedule_bytes("index"),
        )


def test_revision_deduplication_keeps_latest_po() -> None:
    original = dict(_order("jazwares"), filename="PO-100.pdf", version="")
    revision = dict(_order("jazwares"), filename="PO-100 REV2.pdf", version="2")
    selected, report = service._deduplicate([original, revision])
    assert [order["filename"] for order in selected] == ["PO-100 REV2.pdf"]
    assert "忽略 PO-100.pdf" in report[0]


@pytest.mark.parametrize(
    ("customer_code", "text", "expected"),
    [
        (
            "index",
            """INDEX PROMOTIONS
Purchase Order No. PO-INDEX-1
Date 7/1/2026
Ex-Factory 8/1/2026
Qty Units Description Unit Price TOTAL
1234 Product Header
480 pcs. - INDEX Toy $2.00 $960.00
48 pcs/carton
SubTotal""",
            ("PO-INDEX-1", "1234", 480),
        ),
        (
            "jazwares",
            """JAZWARES
07/01/2026 JAZ12345
PO Rev. 2
60 DAYS ROD 08/01/2026 SC-JAZ-1 SEA
ABC123 123-456 JAZWARES Toy 480 48 2.00 960.00
****""",
            ("JAZ12345", "ABC123", 480),
        ),
        (
            "maxx",
            """MAXX MARKETING
P.O. NO. MAXX-PO-1
P.O. DATE 01/07/2026
S.C. NO. SC-MAXX-1
DELIVERY 01/09/2026
MX-100 MAXX Toy 100 PCS 1.50 150.00""",
            ("MAXX-PO-1", "MX-100", 100),
        ),
        (
            "strottman",
            """STROTTMAN
Date Purchase Order # Terms
07/01/2026 PO123 NET
1234-1-F1 6,000 Case $125.00 $750,000.00
STROTTMAN Toy
Special Instructions 750K
01-Aug-2026 15-Aug-2026""",
            ("PO123", "1234-1", 750000),
        ),
        (
            "jp",
            """华康车衣
采购单编号：JP-100
日 期：2026年7月1日
2026年8月1日前交货
4L 布标
12345 车衣 1000 PCS""",
            ("JP-100", "12345", 1000),
        ),
    ],
)
def test_legacy_parser_recognizes_each_huakang_c_customer(
    customer_code: str,
    text: str,
    expected: tuple[str, str, int],
) -> None:
    parsed = service.huakang_po_parser.HuakangPOParser().parse_text(text)
    assert parsed["customer_code"] == service.get_huakang_c_customer_mapping(customer_code).legacy_code
    assert parsed["po_number"] == expected[0]
    assert parsed["lines"][0]["item_code"] == expected[1]
    assert parsed["lines"][0]["qty"] == expected[2]
    if customer_code == "strottman":
        assert parsed["lines"][0]["carton_qty"] == 6000
        assert parsed["lines"][0]["pcs_per_carton"] == 125
        assert parsed["ship_dates"] == ["2026-08-01", "2026-08-15"]


def test_maxx_parser_extracts_project_brand_and_delivery_term() -> None:
    parsed = service.huakang_po_parser.HuakangPOParser().parse_text(
        """MAXX MARKETING
P.O. NO. PO-HK-260222 (REVISION NO. 0)
P.O. DATE 14/4/2026
TERMS OF DELIVERY FCA HK
DELIVERY 6/7/2026
S.C. NO. SC-HK-260349
Project# 260012 Pizza Hut x Toy Story 2026
ITEM NO. ITEM DESCRIPTION QUANTITY U/M USD UNIT PRICE USD AMOUNT
260012-001 Pizza Hut x Toy Story 2026 Cushions 6000 Pieces 2.2700 13,620.00"""
    )
    assert parsed["project_no"] == "260012"
    assert parsed["project_name"] == "Pizza Hut x Toy Story 2026"
    assert parsed["delivery_term"] == "FCA HK"


@pytest.mark.parametrize(
    ("customer_code", "contract_column", "item_column", "quantity_column", "data_row"),
    [
        ("index", 5, 9, 11, 5),
        ("jazwares", 5, 9, 11, 4),
        ("maxx", 5, 9, 11, 5),
        ("strottman", 5, 6, 10, 12),
        ("jp", 4, 5, 7, 4),
    ],
)
def test_preview_and_export_follow_each_customer_schedule_profile(
    monkeypatch,
    customer_code: str,
    contract_column: int,
    item_column: int,
    quantity_column: int,
    data_row: int,
) -> None:
    order = _order(customer_code)
    monkeypatch.setattr(
        service.huakang_po_parser.HuakangPOParser,
        "parse",
        lambda _self, _path: deepcopy(order),
    )
    schedule = _schedule_bytes(customer_code)
    preview = service.create_huakang_c_customer_preview(
        customer_code=customer_code,
        factory_id="huakang-c",
        received_date="2026-08-03",
        po_files=[(f"{customer_code}.pdf", b"pdf")],
        schedule_file_name=f"{customer_code}排期.xlsx",
        schedule_content=schedule,
    )
    row = preview["rows"][0]
    assert preview["factory_id"] == "huakang-c"
    assert row["received_date"] == "2026-08-03"
    assert row["contract_no"] == "SC-100"
    assert row["product_no"] == "ITEM-1"
    assert row["quantity"] == "120"
    assert row["unit_price_hkd"] == "15.5"
    assert row["amount_hkd"] == "1860"
    assert row["product_name_zh" if customer_code == "jp" else "product_name_en"] == "排期标准品名"
    if customer_code == "jp":
        assert row["status"] == "warning"
        assert any("尚无真实PO" in issue["message"] for issue in row["issues"])

    output, file_name, exported_preview = service.export_huakang_c_customer_schedule(
        customer_code=customer_code,
        factory_id="huakang-c",
        received_date="2026-08-03",
        po_files=[(f"{customer_code}.pdf", b"pdf")],
        schedule_file_name=f"{customer_code}排期.xlsx",
        schedule_content=schedule,
    )
    workbook = openpyxl.load_workbook(BytesIO(output), data_only=False)
    try:
        profile = service.huakang_schedule.PROFILES[
            service.get_huakang_c_customer_mapping(customer_code).legacy_code
        ]
        assert workbook.sheetnames == [profile.sheets[0]]
        worksheet = workbook[profile.sheets[0]]
        assert worksheet.cell(profile.style_row, profile.key_cols[0]).value == "OLD-CONTRACT"
        assert worksheet.cell(data_row, contract_column).value == "SC-100"
        assert worksheet.cell(data_row, item_column).value == "ITEM-1"
        assert worksheet.cell(data_row, quantity_column).value == 120
        if customer_code == "index":
            assert worksheet.cell(data_row, 1).value.strftime("%Y-%m-%d") == "2026-08-03"
            assert worksheet.cell(data_row, 4).value == "PO-100"
            assert worksheet.cell(data_row, 6).value == "INDEX"
            assert worksheet.cell(data_row, 8).value == "ITEM-1"
            assert worksheet.cell(data_row, 10).value == "排期标准品名"
            assert worksheet.cell(data_row, 24).value == 12
            assert worksheet.cell(data_row, 25).value == f'=IFERROR(K{data_row}/X{data_row},"")'
            assert worksheet.cell(data_row, 41).value.strftime("%Y-%m-%d") == "2026-09-01"
            assert worksheet.cell(data_row, 42).value == "=2*7.75"
            assert worksheet.cell(data_row, 43).value == f"=AP{data_row}*K{data_row}"
            assert worksheet.cell(data_row, 44).value.strftime("%Y-%m-%d") == "2026-09-01"
            assert worksheet.cell(data_row, 41).number_format == "yyyy/m/d;@"
            expected_hkd_format = '[$HK$-C04]#,##0.00;\\-[$HK$-C04]#,##0.00'
            assert worksheet.cell(data_row, 42).number_format == expected_hkd_format
            assert worksheet.cell(data_row, 43).number_format == expected_hkd_format
            assert worksheet.cell(data_row, 44).number_format == "yyyy/m/d;@"
    finally:
        workbook.close()
    assert file_name == preview["output_file_name"]
    assert exported_preview["summary"]["total"] == 1


def test_invalid_date_styled_history_value_is_preserved_during_full_export(monkeypatch) -> None:
    order = _order("jazwares")
    monkeypatch.setattr(
        service.huakang_po_parser.HuakangPOParser,
        "parse",
        lambda _self, _path: deepcopy(order),
    )
    schedule = _schedule_bytes("jazwares")
    source = openpyxl.load_workbook(BytesIO(schedule), data_only=False)
    source_ws = source[source.sheetnames[0]]
    source_ws["AS20"] = 300028000000
    source_ws["AS20"].number_format = "yyyy/m/d;@"
    stream = BytesIO()
    source.save(stream)
    source.close()

    output, _file_name, _preview = service.export_huakang_c_customer_schedule(
        customer_code="jazwares",
        factory_id="huakang-c",
        received_date="2026-08-03",
        po_files=[("jazwares.pdf", b"pdf")],
        schedule_file_name="jazwares排期.xlsx",
        schedule_content=stream.getvalue(),
    )
    workbook = openpyxl.load_workbook(BytesIO(output), data_only=False)
    try:
        worksheet = workbook[workbook.sheetnames[0]]
        assert worksheet["AS22"].value == 300028000000
        assert worksheet["AS22"].number_format == "General"
    finally:
        workbook.close()


def test_jazwares_export_appends_small_pos_and_one_parent_total_with_template_styles(
    monkeypatch,
) -> None:
    first = deepcopy(_order("jazwares"))
    first["po_number"] = "PO-100"
    first["contract_no"] = "SC-100"
    first["lines"][0]["item_code"] = "ITEM-1-S"
    first["lines"][0]["qty"] = 120
    second = deepcopy(_order("jazwares"))
    second["po_number"] = "PO-101"
    second["contract_no"] = "SC-101"
    second["lines"][0]["item_code"] = "ITEM-1-M"
    second["lines"][0]["qty"] = 240

    def parse(_self, path):
        return deepcopy(second if "second" in Path(path).name else first)

    monkeypatch.setattr(
        service.huakang_po_parser.HuakangPOParser,
        "parse",
        parse,
    )
    schedule = _schedule_bytes("jazwares")
    output, _file_name, _preview = service.export_huakang_c_customer_schedule(
        customer_code="jazwares",
        factory_id="huakang-c",
        received_date="2026-08-12",
        po_files=[("first.pdf", b"one"), ("second.pdf", b"two")],
        schedule_file_name="JAZWARES排期.xlsx",
        schedule_content=schedule,
    )

    workbook = openpyxl.load_workbook(BytesIO(output), data_only=False)
    try:
        worksheet = workbook[service.huakang_schedule.PROFILES["jazwares"].sheets[0]]
        assert worksheet["D4"].value == "PO-100"
        assert worksheet["D5"].value == "PO-101"
        assert worksheet["H4"].value == "ITEM-1"
        assert worksheet["I4"].value == "ITEM-1-S"
        assert worksheet["I5"].value == "ITEM-1-M"
        assert worksheet["J6"].value == "ITEM-1 合计"
        assert worksheet["K6"].value == "=SUM(K4:K5)"
        assert worksheet["Y4"].value == '=IFERROR(K4/12,"")'
        assert worksheet["AP4"].value == "=2*7.75"
        assert worksheet["AQ4"].value == "=AP4*K4"
        assert worksheet["AR4"].value is None
        assert worksheet["A4"].number_format == worksheet["A2"].number_format
        assert worksheet["D4"]._style == worksheet["D2"]._style
        assert worksheet["J4"]._style == worksheet["J2"]._style
        assert worksheet["J6"]._style == worksheet["J3"]._style
        assert worksheet["K6"]._style == worksheet["K3"]._style
        assert worksheet.row_dimensions[4].height == worksheet.row_dimensions[2].height
        assert worksheet.row_dimensions[6].height == worksheet.row_dimensions[3].height
        assert workbook.calculation.fullCalcOnLoad is True
        assert workbook.calculation.forceFullCalc is True
    finally:
        workbook.close()


def test_maxx_export_uses_headers_and_rows_two_to_eight_group_pattern(
    monkeypatch,
) -> None:
    order = deepcopy(_order("maxx"))
    order.update({
        "po_number": "PO-HK-260222",
        "contract_no": "SC-HK-260349",
        "project_no": "260012",
        "project_name": "Pizza Hut x Toy Story 2026",
        "delivery_term": "FCA HK",
        "ship_to": "",
        "ship_date": "2026-07-06",
        "ship_dates": ["2026-07-06"],
    })
    order["lines"] = [
        {
            **deepcopy(order["lines"][0]),
            "item_code": "260012-001",
            "description": "Pizza Hut x Toy Story 2026 Cushions",
            "qty": 6000,
            "pcs_per_carton": 0,
            "carton_qty": 0,
            "unit_price_usd": 2.27,
            "amount_usd": 13620.0,
        },
        {
            **deepcopy(order["lines"][0]),
            "item_code": "260012-003",
            "description": "Pizza Hut x Toy Story 2026 Tote bags",
            "qty": 6000,
            "pcs_per_carton": 0,
            "carton_qty": 0,
            "unit_price_usd": 2.08,
            "amount_usd": 12480.0,
        },
    ]
    monkeypatch.setattr(
        service.huakang_po_parser.HuakangPOParser,
        "parse",
        lambda _self, _path: deepcopy(order),
    )
    schedule = _schedule_bytes("maxx")
    source = openpyxl.load_workbook(BytesIO(schedule), data_only=False)
    output, _file_name, _preview = service.export_huakang_c_customer_schedule(
        customer_code="maxx",
        factory_id="huakang-c",
        received_date="2026-08-12",
        po_files=[("PO-HK-260222 - signed.pdf", b"pdf")],
        schedule_file_name="MAXX 2026 生产排期表.xlsx",
        schedule_content=schedule,
    )
    workbook = openpyxl.load_workbook(BytesIO(output), data_only=False)
    try:
        source_ws = source["排期"]
        worksheet = workbook["排期"]
        assert worksheet["K2"].value == "=120-120"
        assert worksheet["K3"].value == "=SUM(K2:K2)"
        assert worksheet["A5"].value == "8月12号"
        assert worksheet["D5"].value == "PO-HK-260222"
        assert worksheet["E5"].value == "SC-HK-260349"
        assert worksheet["F5"].value == "MAXX"
        assert worksheet["G5"].value == "Pizza Hut x Toy Story 2026"
        assert worksheet["H5"].value == "260012"
        assert worksheet["I5"].value == "260012-001"
        assert worksheet["I6"].value == "260012-003"
        assert worksheet["K5"].value == 6000
        assert worksheet["K6"].value == 6000
        assert worksheet["L5"].value is None
        assert worksheet["L6"].value is None
        assert worksheet["Y5"].value is None
        assert worksheet["AP5"].value.strftime("%Y-%m-%d") == "2026-07-06"
        assert worksheet["AQ5"].value == "=2.27*7.75"
        assert worksheet["AR5"].value == "=+AQ5*K5"
        assert worksheet["AT5"].value == "FCA HK"
        assert worksheet["J7"].value == "合计"
        assert worksheet["K7"].value == "=SUM(K5:K6)"
        assert worksheet["J8"].value is None
        assert worksheet["K8"].value is None
        assert worksheet["D5"]._style == source_ws["D2"]._style
        assert worksheet["A5"]._style == source_ws["A2"]._style
        assert worksheet["AQ5"]._style == source_ws["AQ2"]._style
        assert worksheet["J7"]._style == source_ws["J3"]._style
        assert worksheet["K7"]._style == source_ws["K3"]._style
        assert worksheet.row_dimensions[5].height == source_ws.row_dimensions[2].height
        assert worksheet.row_dimensions[7].height == source_ws.row_dimensions[3].height
        assert workbook.calculation.fullCalcOnLoad is True
        assert workbook.calculation.forceFullCalc is True
    finally:
        source.close()
        workbook.close()


def test_strottman_export_uses_row_three_fields_and_rows_four_to_ten_layout(
    monkeypatch,
) -> None:
    order = deepcopy(_order("strottman"))
    order.update({
        "po_number": "PO953",
        "contract_no": "PO953",
        "ship_date": "2026-08-21",
        "ship_dates": ["2026-08-21", "2026-11-19"],
        "contact": "朱江",
    })
    order["lines"] = [{
        **deepcopy(order["lines"][0]),
        "item_code": "0229-672",
        "description": "6in Standing Cow",
        "qty": 750000,
        "pcs_per_carton": 125,
        "carton_qty": 6000,
        "unit_price_usd": 1.275,
        "amount_usd": 956250.0,
    }]
    monkeypatch.setattr(
        service.huakang_po_parser.HuakangPOParser,
        "parse",
        lambda _self, _path: deepcopy(order),
    )
    schedule = _schedule_bytes("strottman")
    source = openpyxl.load_workbook(BytesIO(schedule), data_only=False)
    output, _file_name, preview = service.export_huakang_c_customer_schedule(
        customer_code="strottman",
        factory_id="huakang-c",
        received_date="2026-08-12",
        po_files=[("0229-672 6in Standing Plush PO953_RR.pdf", b"pdf")],
        schedule_file_name="Strottman_排期表.xlsx",
        schedule_content=schedule,
    )
    workbook = openpyxl.load_workbook(BytesIO(output), data_only=False)
    try:
        source_ws = source["建文客排期表"]
        worksheet = workbook["建文客排期表"]
        assert preview["rows"][0]["item_sheet_name"] == "建文客排期表"
        assert "A11:C11" in worksheet.merged_cells
        assert "D11:K11" in worksheet.merged_cells
        assert worksheet["A11"].value == "0229-672"
        assert worksheet["D11"].value == "6in Standing Cow"
        assert worksheet["A12"].value == "正常备料"
        assert worksheet["B12"].value == "正常生产"
        assert worksheet["C12"].value == "0229-672"
        assert worksheet["D12"].value.strftime("%Y-%m-%d") == "2026-08-12"
        assert worksheet["E12"].value == 953
        assert worksheet["F12"].value == "0229-672"
        assert worksheet["G12"].value == "0229-672"
        assert worksheet["H12"].value == "6in Standing Cow"
        assert worksheet["J12"].value == 750000
        assert worksheet["M12"].value.strftime("%Y-%m-%d") == "2026-08-21"
        assert worksheet["M12"].number_format == 'm"月"d"日"'
        assert worksheet["R12"].value == "Hong Kong"
        assert worksheet["S12"].value == "STROTTMAN"
        assert worksheet["U12"].value == "朱江"
        assert worksheet["W12"].value == 125
        assert worksheet["X12"].value == '=IFERROR(J12/W12,"")'
        assert worksheet["AA12"].value == 1.275
        assert worksheet["AB12"].value == "=AA12*7.75"
        assert worksheet["AC12"].value == "=AA12*J12"
        assert worksheet["AD12"].value == "=AB12*J12"
        assert worksheet["AF12"].value == '=IF(AE12="","",AE12*J12)'
        assert "2026-08-21 / 2026-11-19" in worksheet["AG12"].value
        assert worksheet["AH12"].value.strftime("%Y-%m-%d") == "2026-08-12"
        assert worksheet["A11"]._style == source_ws["A9"]._style
        assert worksheet["D11"]._style == source_ws["D9"]._style
        style_mismatches = [
            col_no
            for col_no in range(1, 47)
            if col_no != 13
            and worksheet.cell(12, col_no)._style != source_ws.cell(10, col_no)._style
        ]
        assert style_mismatches == []
        assert worksheet.row_dimensions[11].height == source_ws.row_dimensions[9].height
        assert worksheet.row_dimensions[12].height == source_ws.row_dimensions[10].height
        assert workbook.calculation.fullCalcOnLoad is True
        assert workbook.calculation.forceFullCalc is True
    finally:
        source.close()
        workbook.close()


def test_cross_customer_po_is_rejected(monkeypatch) -> None:
    wrong_order = _order("maxx")
    monkeypatch.setattr(
        service.huakang_po_parser.HuakangPOParser,
        "parse",
        lambda _self, _path: deepcopy(wrong_order),
    )
    with pytest.raises(service.HuakangCCustomerOrderError, match="当前入口只处理INDEX"):
        service.create_huakang_c_customer_preview(
            customer_code="index",
            factory_id="huakang-c",
            received_date="2026-08-03",
            po_files=[("MAXX.pdf", b"pdf")],
            schedule_file_name="INDEX排期.xlsx",
            schedule_content=_schedule_bytes("index"),
        )
