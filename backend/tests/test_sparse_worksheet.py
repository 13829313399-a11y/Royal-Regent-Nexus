from io import BytesIO

import openpyxl
import pytest
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.sparse_worksheet import insert_rows
from app.services.huakang_a_order_legacy.unified_schedule import _insert_rows


def test_sparse_insert_retains_distant_content_without_allocating_empty_rectangle():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet['A4'] = '现有订单'
    sheet['A5'] = '取消单'
    sheet['D9'] = '已走货订单'
    sheet['AK100000'] = '用户保留备注'
    sheet['AK100000'].fill = PatternFill('solid', fgColor='FF0000')
    original = len(sheet._cells)
    insert_rows(sheet, 5, 3)
    assert len(sheet._cells) == original
    assert sheet['A4'].value == '现有订单'
    assert sheet['A8'].value == '取消单'
    assert sheet['D12'].value == '已走货订单'
    assert sheet['AK100003'].value == '用户保留备注'
    assert sheet['AK100003'].fill.fgColor.rgb == '00FF0000'
    output = BytesIO(); workbook.save(output)
    restored = openpyxl.load_workbook(BytesIO(output.getvalue()))
    assert restored.active['AK100003'].value == '用户保留备注'
    assert len(restored.active._cells) < 10


def test_regional_insert_preserves_formulas_merges_validation_and_row_dimensions():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'ITEM表'
    sheet['A4'] = 7
    sheet['A5'] = '取消单'
    sheet['B8'] = '=A4+A8'
    sheet['AK100000'] = '保留'
    sheet.merge_cells('C8:D8')
    sheet.row_dimensions[8].height = 42
    validation = DataValidation(type='whole', formula1=1, formula2=100)
    validation.add('A8:A10'); sheet.add_data_validation(validation)
    _insert_rows(workbook, sheet, 5, 2)
    assert len(sheet._cells) < 20
    assert sheet['B10'].value == '=A4+A10'
    assert 'C10:D10' in sheet.merged_cells
    assert sheet.row_dimensions[10].height == 42
    assert str(validation.sqref) == 'A10:A12'
    assert sheet['AK100002'].value == '保留'


def test_excel_row_limit_rejects_without_mutating_original():
    sheet = openpyxl.Workbook().active
    sheet['A1048576'] = '保留'
    with pytest.raises(ValueError, match='行数上限'):
        insert_rows(sheet, 4)
    assert sheet['A1048576'].value == '保留'
    assert len(sheet._cells) == 1
