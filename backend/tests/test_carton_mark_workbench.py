import io
import json

import pytest
from fastapi import HTTPException
from PIL import Image
from pypdf import PdfReader

from app.services import carton_mark_workbench as workbench
from app.services import carton_mark_flexible as flow
from app.services import carton_supplier_mark_ai as ai
from app.services import carton_mark_ai_pilot as pilot
from test_carton_mark_flexible import table_workbook,table_plan


def setup(monkeypatch):
    content=table_workbook();original=dict(excel_file_name="original.xlsx",excel_bytes=content)
    source=flow.read("original.xlsx",content);calls=[]
    monkeypatch.setattr(ai,"configuration",lambda:({},pilot.load_template(),b"reference","signature"))
    monkeypatch.setattr(workbench.carton_mark_images,"collect",lambda *args,**kwargs:([],[],[]))
    image=io.BytesIO();Image.new("RGB",(200,100),"black").save(image,format="PNG")
    monkeypatch.setattr(flow,"crop_reference",lambda *args:image.getvalue())
    def provider(*args,**kwargs):calls.append(args);return json.dumps(table_plan(source)),{}
    monkeypatch.setattr(pilot,"request_plan",provider)
    return original,calls


def test_read_once_then_manual_weight_and_note_review_are_rendered_without_another_provider_call(monkeypatch):
    original,calls=setup(monkeypatch);read=workbench.read(original,"owner")
    assert len(calls)==1 and len(read["records"])==2
    field=next(f for f in read["records"][0]["panels"][1]["fields"] if f["label"]=="NW:")
    pdf,pages,warnings=workbench.render(original,"owner",read["payload"],read["signature"],
        {field["id"]:"2.5 KGS"},[read["notes"][0]["id"]],"A3","landscape")
    assert len(calls)==1 and pages>=3
    text="\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "2.5 KGS" in text and "页面补充" in text and "页面已明确核对" in text
    assert any("人工补充" in w for w in warnings)
    workbench.render(original,"owner",read["payload"],read["signature"],{},[],"A3","portrait")
    assert len(calls)==1


@pytest.mark.parametrize("kind",["owner","file","payload","expired","invalid_edit","po_edit","unknown_note"])
def test_review_cannot_be_reused_for_another_owner_source_or_edited_plan(monkeypatch,kind):
    original,calls=setup(monkeypatch);read=workbench.read(original,"owner")
    owner="owner";payload=read["payload"];signature=read["signature"];edits={};reviewed=[]
    if kind=="owner":owner="different-owner"
    if kind=="file":original={**original,"excel_bytes":original["excel_bytes"]+b"changed"}
    if kind=="payload":payload=payload.replace("PO-SYNTHETIC-A","FORGED-PO")
    if kind=="expired":monkeypatch.setattr(workbench.time,"time",lambda:9_999_999_999)
    if kind=="invalid_edit":edits={"r999p0f0":"test"}
    if kind=="po_edit":
        field=next(f for f in read["records"][0]["panels"][0]["fields"] if f["locked"]);edits={field["id"]:"DIFFERENT-PO"}
    if kind=="unknown_note":reviewed=["N999"]
    with pytest.raises(HTTPException):workbench.render(original,owner,payload,signature,edits,reviewed,"A3","landscape")
    assert len(calls)==1


def test_local_image_reference_text_cannot_change_native_business_values_or_barcode_dimensions(monkeypatch):
    original,_=setup(monkeypatch)
    source=flow.read("original.xlsx",original["excel_bytes"])
    source["notes"].append(pilot.Evidence(id="I1",kind="image_reference_text",location="原图!image:1",text="旧示例条码尺寸900mm x 900mm；旧PO123456"))
    assert flow.constraints(source)["height"]==32
    result=table_plan(source)
    result["records"][0]["panels"][0]["fields"][0]["parts"]=[dict(source_id="I1",value="123456")]
    with pytest.raises(pilot.PilotError):flow.validate(json.dumps(result),source)


def test_source_only_fallback_and_manual_addition_keep_native_values_traceable(monkeypatch):
    original,calls=setup(monkeypatch)
    def incomplete(*a,**k):calls.append(a);raise pilot.PilotError('模型响应不完整')
    monkeypatch.setattr(pilot,'request_plan',incomplete)
    read=workbench.read(original,'owner');assert read['mode']=='needs_mapping' and len(read['records'])==2
    pdf,_,_=workbench.render(original,'owner',read['payload'],read['signature'],{},[],'A3','portrait',
        [dict(record=0,face='side',label='NW',value='2.5 KGS')],[])
    text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert '2.5 KGS' in text and '用户填写' in text and 'AI排版建议未能通过核验' in text
    assert len(calls)==1


def test_source_only_mark_pairs_native_labels_values_and_units_without_truncating_faces(monkeypatch):
    from test_carton_mark_flexible import workbook
    from app.services.carton_mark_native_plan import build
    original,calls=setup(monkeypatch);original['excel_bytes']=workbook()
    src=flow.read('original.xlsx',original['excel_bytes'],allow_unresolved=True)
    plan=build(src,'synthetic incomplete response')
    assert plan.orientation=='landscape'
    assert [p.face for p in plan.records[0].panels]==['front','side']
    front,side=plan.records[0].panels
    field=next(f for f in front.fields if f.label and f.label.value=='ORDER NUMBER:')
    assert flow.field_value(field.parts)=='PO-DEMO-0001'
    qty=next(f for f in front.fields if f.label and f.label.value=='QTY:')
    assert flow.field_value(qty.parts)=='6 PCS/CTN'
    assert any(f.label and f.label.value=='CARTON NO:' for f in front.fields)
    assert front.barcodes and side.barcodes
    assert all('mark (two sides)' not in part.value.lower() for p in (front,side) for f in p.fields for part in f.parts)
    def incomplete(*args,**kwargs):calls.append(args);raise pilot.PilotError('synthetic incomplete response')
    monkeypatch.setattr(pilot,'request_plan',incomplete)
    read=workbench.read(original,'owner')
    assert read['mode']=='needs_mapping' and read['orientation']=='landscape'
    pdf,_,_=workbench.render(original,'owner',read['payload'],read['signature'],{},[],read['paper'],read['orientation'],include_sources=False)
    page=PdfReader(io.BytesIO(pdf)).pages[0];lines=[]
    page.extract_text(visitor_text=lambda text,cm,tm,font,size:lines.append(text.strip()) if text.strip() else None)
    text=page.extract_text()
    assert page.mediabox.width>page.mediabox.height
    assert 'ORDER NUMBER: PO-DEMO-0001' in lines
    assert 'ORDER NUMBER: PO-DEMO-0001' in text
    assert 'QTY: 6 PCS/CTN' in text
    assert len(calls)==1


def test_compact_pointer_plan_materializes_unchanged_native_text_and_rejects_unknown_ids(monkeypatch):
    from app.services.carton_mark_compact_plan import expand
    original,_=setup(monkeypatch);source=flow.read('original.xlsx',original['excel_bytes']);plan=table_plan(source)
    for record in plan['records']:
        for panel in record['panels']:
            for field in panel['fields']:
                field['source_ids']=[p['source_id'] for p in field.pop('parts')]
                label=field.pop('label');field['label_id']=label['source_id'] if label else None
            for bar in panel['barcodes']:bar['source_id']=bar.pop('source')['source_id']
    raw=expand(json.dumps(plan),source);assert '000123' in raw
    valid=flow.validate(raw,source,review_mode=True);assert len(valid.records)==2
    plan['records'][0]['panels'][0]['fields'][0]['source_ids']=['E999999']
    with pytest.raises(pilot.PilotError):expand(json.dumps(plan),source)


def test_manual_native_barcode_is_checked_and_another_po_source_is_refused(monkeypatch):
    original,calls=setup(monkeypatch);read=workbench.read(original,'owner')
    source=flow.read('original.xlsx',original['excel_bytes'])
    good=next(c for c in source['cells'] if c.location.endswith('!F24'))
    wrong=next(c for c in source['cells'] if c.location.endswith('!F25'))
    args=(original,'owner',read['payload'],read['signature'],{},[],'A3','portrait')
    pdf,_,_=workbench.render(*args,[dict(record=0,face='side',source_id=good.id,as_barcode=True)],[])
    assert '4006381333931' in '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    with pytest.raises(HTTPException):workbench.render(*args,[dict(record=0,face='side',source_id=wrong.id,as_barcode=True)],[])
    assert len(calls)==1


def test_each_table_uses_its_own_header_and_prefers_customer_po_to_generic_po():
    from openpyxl import Workbook
    from app.services.carton_mark_native_plan import build
    b=Workbook();s=b.active
    for row,dc in ((1,'DC#881'),(5,'DC#891')):
        for col,text in enumerate(('PO NO.','Customer PO','ITEM NO.','Case Pack','QTY',dc),1):s.cell(row,col,text)
        for col,text in enumerate(('CONTRACT','PO-ONE' if row==1 else 'PO-TWO','000123','6','12','2'),1):s.cell(row+1,col,text)
    out=io.BytesIO();b.save(out);b.close()
    source=flow.read('two-tables.xlsx',out.getvalue());assert len(source['detail_rows'])==2
    plan=build(source,'test')
    assert plan.records[0].panels[0].fields[-1].label.value=='DC#881'
    assert plan.records[1].panels[0].fields[-1].label.value=='DC#891'


def test_wrong_native_gtin_is_preserved_without_a_barcode_graph_and_cannot_bypass_checks_with_code128():
    from openpyxl import load_workbook
    from test_carton_mark_flexible import workbook,plan
    b=load_workbook(io.BytesIO(workbook()));b.active['C11']='4006381333932'
    out=io.BytesIO();b.save(out);b.close();source=flow.read('bad-gtin.xlsx',out.getvalue())
    raw=plan(source);raw['records'][0]['panels'][0]['barcodes'][0]['kind']='Code128'
    valid=flow.validate(json.dumps(raw),source,review_mode=True)
    assert not valid.records[0].panels[0].barcodes
    assert any(p.value=='4006381333932' for f in valid.records[0].panels[0].fields for p in f.parts)
    assert any('未生成条码图' in w for w in valid._warnings)


def test_explicit_graphic_choices_use_controlled_artwork_and_are_traced_without_api_recall(monkeypatch):
    original,calls=setup(monkeypatch);read=workbench.read(original,'owner')
    result=workbench.render(original,'owner',read['payload'],read['signature'],{},[],'A3','portrait',[],[],
        [dict(record=0,face='side',logo=True,symbols=['keep_dry'])])
    text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(result[0])).pages)
    assert '防潮' in text and '图案' in text and len(calls)==1
    with pytest.raises(HTTPException):
        workbench.render(original,'owner',read['payload'],read['signature'],{},[],'A3','portrait',[],[],
            [dict(record=0,face='side',logo=True,symbols=['external-svg-code'])])


def test_instruction_sentences_do_not_become_duplicate_face_headers():
    from openpyxl import load_workbook
    from test_carton_mark_flexible import workbook
    book=load_workbook(io.BytesIO(workbook()))
    book.active['D20']='正唛的外箱条码下需要加地址'
    book.active['D21']='侧唛内容需要核对重量'
    stream=io.BytesIO();book.save(stream);book.close()
    source=flow.read('face-notes.xlsx',stream.getvalue())
    identities={c.location.split('!')[1]:c.id for c in source['cells']}
    assert source['ownership'][identities['C6']]=='front'
    assert source['ownership'][identities['G6']]=='side'


def test_six_large_native_barcodes_get_full_width_faces_without_shrinking():
    from test_carton_mark_flexible import workbook,plan
    source=flow.read('multiple-barcodes.xlsx',workbook())
    parsed=flow.Plan.model_validate(plan(source))
    for panel in parsed.records[0].panels:
        panel.logo=False;panel.fields=[];panel.symbols=[]
    side=parsed.records[0].panels[1]
    original=side.barcodes[0]
    side.barcodes=[original.model_copy(deep=True) for _ in range(6)]
    pdf,pages,warnings=flow.render(parsed,source)
    assert pages>=3 and all(b.height_mm==60 and b.width_mm==80 for b in side.barcodes)
    assert any('各面独立一页' in w for w in warnings)
    assert any('分列展示' in w for w in warnings)
    assert '4006381333931' in PdfReader(io.BytesIO(pdf)).pages[1].extract_text()


def multi_sheet_notes_workbook():
    from openpyxl import Workbook
    book=Workbook();book.remove(book.active)
    for name in ('AUSYD','AUBNE','AUMEL','AUFRE'):
        sheet=book.create_sheet(name);sheet['A1']='CURRENT-ITEM-'+name
        for row in range(2,46):sheet.cell(row,1,f'备注 {row}: 核对当前纸箱要求')
    out=io.BytesIO();book.save(out);book.close();return out.getvalue()


@pytest.mark.parametrize('provider_fails',[False,True])
def test_multi_sheet_notes_survive_read_all_reviewed_and_rerender_without_repeat_api(monkeypatch,provider_fails):
    content=multi_sheet_notes_workbook();original=dict(excel_file_name='multiple-ports.xlsx',excel_bytes=content)
    source=flow.read(original['excel_file_name'],content,allow_unresolved=True)
    assert len(source['notes'])==176
    assert len(flow.read(original['excel_file_name'],content,allow_unresolved=True,worksheets=['AUSYD'])['notes'])==44
    with pytest.raises(pilot.PilotError,match='备注过多'):flow.read(original['excel_file_name'],content)
    monkeypatch.setattr(ai,'configuration',lambda:({},pilot.load_template(),b'reference','signature'))
    monkeypatch.setattr(workbench.carton_mark_images,'collect',lambda *a,**k:([],[],[]))
    calls=[]
    def provider(*a,**k):
        calls.append(a)
        if provider_fails:raise pilot.PilotError('synthetic incomplete response')
        return json.dumps(dict(template_id='source-flow-v1',paper='A3',orientation='portrait',records=[
            dict(panels=[dict(face='front',copies=2,logo=False,fields=[dict(label_id=None,source_ids=[c.id],role='normal')],barcodes=[],symbols=[])])
            for c in source['cells'] if c.text.startswith('CURRENT-ITEM-')])),{}
    monkeypatch.setattr(pilot,'request_plan',provider)
    read=workbench.read(original,'owner')
    assert read['mode']==('needs_mapping' if provider_fails else 'ai') and len(read['notes'])==176
    selected=[n['id'] for n in read['notes']]
    pdf,_,_=workbench.render(original,'owner',read['payload'],read['signature'],{},selected,'A3','portrait')
    text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert all(n['location'] in text for n in read['notes'])
    assert text.count('页面已明确核对')==176
    workbench.render(original,'owner',read['payload'],read['signature'],{},selected,'A3','landscape')
    assert len(calls)==1


@pytest.mark.parametrize('overflow',['count','characters','single_note'])
def test_extended_note_bounds_fail_explicitly_without_dropping_source_notes(overflow):
    from types import SimpleNamespace
    from app.services.carton_mark_source_notes import collect
    texts=([f'备注 {i}' for i in range(2501)] if overflow=='count' else
           ['备注'+'字'*3998]*31 if overflow=='characters' else ['备注'+'字'*3999])
    doc=SimpleNamespace(items=[SimpleNamespace(location=f'Sheet1!A{i+1}',text=t) for i,t in enumerate(texts)])
    with pytest.raises(pilot.PilotError,match='未调用API'):
        collect('bounded.xls',b'',extracted=doc,extended=True)


def test_clean_preview_and_complete_record_preserve_distinct_content_without_api_recall(monkeypatch):
    original,calls=setup(monkeypatch);read=workbench.read(original,'owner')
    args=(original,'owner',read['payload'],read['signature'],{},[],'A3','landscape')
    preview,pages,_=workbench.render(*args,include_sources=False)
    full,full_pages,_=workbench.render(*args,include_sources=True)
    clean_text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(preview)).pages)
    full_text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(full)).pages)
    assert pages==len(read['records']) and full_pages>pages
    assert 'PO-SYNTHETIC-A' in clean_text and '制版备注与来源' not in clean_text
    assert '制版备注与来源' in full_text and all(n['location'] in full_text for n in read['notes'])
    assert len(calls)==1


def test_model_gets_native_structure_and_only_source_cited_uncertainties_are_shown(monkeypatch):
    from app.services import carton_mark_reading as reading
    original,_=setup(monkeypatch);source=flow.read('original.xlsx',original['excel_bytes'])
    facts=reading.structure(source)
    assert len(facts['current_detail_rows'])==2 and any(s['role_hint']=='reference' for s in facts['worksheets'])
    candidates=reading.provider_evidence(source)
    example_po=next(c for c in source['cells'] if c.location.endswith('!C6'))
    assert next(c for c in candidates if c.id==example_po.id).kind=='template_example_value'
    assert next(c for c in candidates if c.location.endswith('!D9')).text=='PCS/CTN'
    assert source['cells'][next(i for i,c in enumerate(source['cells']) if c.id==example_po.id)].text=='PO-DEMO-0001'
    raw=table_plan(source);note=source['notes'][0]
    raw['uncertainties']=[dict(source_ids=[note.id],message='请核实这项制版要求'),dict(source_ids=['E999999'],message='伪造来源')]
    received=[]
    def provider(entries,*a,**k):received.extend(entries);return json.dumps(raw),{}
    monkeypatch.setattr(pilot,'request_plan',provider)
    read=workbench.read(original,'owner')
    assert read['mode']=='ai' and any(e.id=='S1' and e.kind=='source_structure' for e in received)
    questions=[i for i in read['issues'] if i['kind']=='question']
    assert len(questions)==1 and questions[0]['sources']==[note.location]
    assert any(i['kind']=='missing_value' and i['field_ids'] for i in read['issues'])
    assert all('伪造来源' not in i['message'] for i in read['issues'])
    full,_,_=workbench.render(original,'owner',read['payload'],read['signature'],{},[],'A3','landscape')
    assert '请核实这项制版要求' in '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(full)).pages)
    source['notes'].append(pilot.Evidence(id='Q1',kind='model_question',location='AI待确认!1',text='条码高度最低900mm，请确认'))
    assert flow.constraints(source)['height']==32


def test_empty_model_placeholders_do_not_discard_native_rows_weights_or_required_address():
    source=flow.read('table.xlsx',table_workbook());raw=table_plan(source)
    for record in raw['records']:
        for panel in record['panels']:
            panel['copies']=4
            panel['fields']=[f for f in panel['fields'] if f['role'] not in {'placeholder','address'}]
            panel['fields'].append(dict(label=None,parts=[],role='placeholder'))
    plan=flow.validate(json.dumps(raw),source,review_mode=True)
    assert len(plan.records)==2
    for record in plan.records:
        assert all(p.copies==2 for p in record.panels)
        assert sum(f.role=='placeholder' for p in record.panels for f in p.fields)==2
        addresses=[f for p in record.panels for f in p.fields if f.role=='address']
        assert addresses and addresses[0].position=='after_barcode'
