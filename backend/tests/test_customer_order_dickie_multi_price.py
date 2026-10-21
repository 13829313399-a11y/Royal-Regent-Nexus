"""Synthetic release data exercises per-PO price allocation and OCR handling."""
from copy import deepcopy
from decimal import Decimal

import pytest

from app.services import customer_order_dickie as service


def multi_price_order(second_price="13.75"):
    return service._parse_dickie_ocr_pages([
        """
        Release Order Page 1
        Reference: SC900000101/ 100
        Date of creation: 12.SEP.2030
        Mat. No.: 909090101 Mat. EAN:
        Sample Camper Set, Try Me
        Packing: OPC / 6PC
        Dickie open box
        Quantity Master Contract No. PO Contract No. Unit Price Delivery Date
        120 PC 。 500090001/ 10 300990001/ 10 12.25 HKD 21.FEB.2031
        360 PC 500090002/ 10 300990002/ 10 """ + second_price + """ HKD 21.FEB.2031
        """,
        "This order is for SAMPLE RETAIL. Port of discharge: SAMPLE PORT",
    ], fallback_received_date="2026-10-15")


def preview(parsed):
    return service._preview_row(parsed, None, {}, file_name="SC900000101.pdf", row_index=1)


def test_homepage_keeps_each_po_quantity_price_and_exact_amount():
    parsed = multi_price_order()
    assert parsed.quantity == 480
    assert parsed.po_no == "300990001/10\n300990002/10"
    assert parsed.unit_price_hkd == Decimal("13.375")
    assert parsed.allocations == [("500090001/10", Decimal("120")), ("500090002/10", Decimal("360"))]
    row = preview(parsed)
    assert row["contract_no"] == "500090001/10 120pc\n500090002/10 360pc"
    assert row["price_basis"] == "weighted"
    assert row["amount_hkd"] == "6420"
    assert [item["unit_price_hkd"] for item in row["price_details"]] == ["12.25", "13.75"]
    assert [item["amount_hkd"] for item in row["price_details"]] == ["1470", "4950"]
    assert "加权均价" in row["price_details_text"]


def test_unreadable_second_price_keeps_both_po_rows_and_can_be_corrected():
    row = preview(multi_price_order("om"))
    assert row["quantity"] == "480"
    assert row["unit_price_hkd"] == row["amount_hkd"] == ""
    assert row["price_basis"] == "incomplete"
    assert row["price_details"][1]["unit_price_hkd"] == ""
    issue = next(issue for issue in row["issues"] if issue["code"] == "missing_po_unit_prices")
    assert issue["severity"] == "warning"
    assert issue["edit_field"] == "po_unit_prices"
    service.apply_po_price_overrides(row, "300990002/10=13.75")
    assert row["unit_price_hkd"] == "13.375"
    assert row["amount_hkd"] == "6420"
    assert row["price_basis"] == "weighted"
    updated_issue = next(issue for issue in row["issues"] if issue["code"] == "missing_po_unit_prices")
    assert updated_issue["edit_field"] == "po_unit_prices"
    assert "已人工补录" in updated_issue["message"]


@pytest.mark.parametrize("value", [
    "999999999/10=13.75", "300990002/10=0", "300990002/10=-1",
    "300990002/10=13.75\n300990002/10=52", "300990002/10=13.75\ninvalid",
])
def test_price_overrides_reject_unknown_invalid_and_duplicate_rows_atomically(value):
    row = preview(multi_price_order("om"))
    original = deepcopy(row)
    with pytest.raises(service.CustomerOrderWorkbookError):
        service.apply_po_price_overrides(row, value)
    assert row == original


def test_price_overrides_do_not_reallocate_changed_order_quantity():
    row = preview(multi_price_order("om"))
    row["quantity"] = "500"
    original = deepcopy(row)
    with pytest.raises(service.CustomerOrderWorkbookError, match="数量合计"):
        service.apply_po_price_overrides(row, "300990002/10=13.75")
    assert row == original


def test_ocr_space_inside_po_and_october_zero_are_normalized():
    parsed = service._parse_dickie_ocr_pages([
        """Reference: SC900000102/1300
        Date of creation: 02.0CT.2030
        Mat. No.: 909090102
        Sample Cable Car
        Packing: 0 PC / 12 PC
        96 PC 500090003/ 10 300990 003/ 10 8.50 HKD 14.JUN.2031"""
    ], fallback_received_date="2026-10-15")
    assert parsed.po_no == "300990003/10"
    assert parsed.source_date == "2030-10-02"
    assert parsed.quantity == 96
    assert parsed.unit_price_hkd == Decimal("8.50")


def test_per_po_override_survives_unified_export_and_api_control(monkeypatch):
    from io import BytesIO
    import openpyxl
    from app.api import customer_order as api
    from app.services import customer_order_unified as unified
    from app.services.customer_order_manual import decorate_manual_resolution_policy
    from test_customer_order_huaxing_unified import template, data

    book = template("dickie")
    book["ITEM表"]["H4"] = "90 909 0101"
    book["ITEM表"]["I4"] = "样例产品"
    monkeypatch.setattr(service, "parse_dickie_pdf_orders", lambda *args, **kwargs: [multi_price_order("om")])
    args = dict(customer_code="dickie", factory_id="huaxing", received_date="2026-10-15",
                po_files=[("release.pdf", b"%PDF-1.7")], schedule_file_name="schedule.xlsx", schedule_content=data(book))
    initial = decorate_manual_resolution_policy(unified.create_unified_customer_preview(**args))
    row, = initial["rows"]
    issue = next(i for i in row["issues"] if i["code"] == "missing_po_unit_prices")
    override = dict(row_id=row["id"], issue_key=issue["skip_key"], field="po_unit_prices", value="300990002/10=13.75")
    assert api._actual_manual_overrides(initial, [override]) == [override]
    output, filename, exported = unified.export_unified_customer_schedule(**args, manual_overrides=[override])
    assert exported["rows"][0]["amount_hkd"] == "6420"
    assert exported["rows"][0]["unit_price_hkd"] == "13.375"
    written = openpyxl.load_workbook(BytesIO(output))
    target = next(r for r in range(4, written["接单表"].max_row+1) if written["接单表"].cell(r,4).value == row["reference_no"])
    assert written["接单表"].cell(target,14).value == pytest.approx(13.375)
    monkeypatch.setattr(api, "_record_export_audit", lambda *args, **kwargs: "audit")
    _, audit = api._complete_export_control(None, None, preview=exported, received_date=args["received_date"],
        submitted_fingerprint=api._preview_fingerprint(initial,args["received_date"]),
        requested_issue_keys=set(), requested_manual_overrides=[override], confirmation_reason="已按原单逐PO核对价格",
        output=output, file_name=filename)
    assert audit == "audit"


def test_legacy_export_keeps_per_po_prices_and_accepts_api_audit(monkeypatch):
    from io import BytesIO
    import openpyxl
    from app.api import customer_order as api
    from app.services.customer_order_buzzbee import _decrypt_schedule
    from app.services.customer_order_manual import decorate_manual_resolution_policy
    from test_customer_order_dickie import build_dickie_schedule

    monkeypatch.setattr(service, "parse_dickie_pdf_orders", lambda *args, **kwargs: [multi_price_order("om")])
    args = dict(factory_id="huaxing", received_date="2026-10-15", po_files=[("release.pdf",b"%PDF-1.7")],
                schedule_file_name="legacy.xlsx",schedule_content=build_dickie_schedule())
    initial = decorate_manual_resolution_policy(service.create_dickie_batch_preview(**args))
    row, = initial["rows"]
    issue = next(i for i in row["issues"] if i["code"] == "missing_po_unit_prices")
    assert issue["can_skip"] and issue["skip_key"]
    override = dict(row_id=row["id"],issue_key=issue["skip_key"],field="po_unit_prices",value="300990002/10=13.75")
    assert api._actual_manual_overrides(initial,[override]) == [override]
    output,filename,exported=service.export_dickie_batch_schedule(**args,skipped_issue_keys={issue["skip_key"]},manual_overrides=[override])
    assert exported["rows"][0]["unit_price_hkd"] == "13.375"
    assert exported["rows"][0]["amount_hkd"] == "6420"
    plain,_ = _decrypt_schedule(output)
    written=openpyxl.load_workbook(BytesIO(plain))
    for sheet_name in ("Iteam表","接单表","正单评审表"):
        sheet=written[sheet_name]
        ref_col=3 if sheet_name=="Iteam表" else 4
        target=next(r for r in range(4,sheet.max_row+1) if sheet.cell(r,ref_col).value==row["reference_no"])
        price_col=next(c.column for c in sheet[3] if c.value=="PO单价明细")
        basis_col=next(c.column for c in sheet[3] if c.value=="单价口径")
        assert "300990001/10：120 × 12.25" in sheet.cell(target,price_col).value
        assert "300990002/10：360 × 13.75" in sheet.cell(target,price_col).value
        assert sheet.cell(target,basis_col).value=="数量加权均价"
        if sheet_name=="正单评审表":
            assert sheet.cell(target,19).value==pytest.approx(13.375)
    monkeypatch.setattr(api,"_record_export_audit",lambda *args,**kwargs:"audit")
    _,audit=api._complete_export_control(None,None,preview=exported,received_date=args["received_date"],
        submitted_fingerprint=api._preview_fingerprint(initial,args["received_date"]),
        requested_issue_keys={issue["skip_key"]},requested_manual_overrides=[override],confirmation_reason="已按原单核对",
        output=output,file_name=filename)
    assert audit=="audit"
