import sys
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.shared import Inches, Pt
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.styles import Border, Font, PatternFill, Side
from PIL import Image as PillowImage


BACKEND_DIR = Path(__file__).resolve().parents[1]

if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))


def _png_bytes() -> bytes:
    output = BytesIO()
    PillowImage.new("RGB", (24, 16), color=(93, 64, 211)).save(output, format="PNG")
    return output.getvalue()


def _fake_translator(texts, direction):
    assert direction in {"zh_to_en", "en_to_zh"}
    mappings = {
        "采购订单": "Purchase Order",
        "数量": "Quantity",
        "页眉说明": "Header note",
        "页脚说明": "Footer note",
        "Purchase Order": "采购订单",
        "Quantity": "数量",
    }
    return [mappings.get(text, f"translated:{text}") for text in texts]


def _archive_media(content: bytes, prefix: str) -> dict[str, bytes]:
    with ZipFile(BytesIO(content)) as archive:
        return {
            name: archive.read(name)
            for name in archive.namelist()
            if name.startswith(prefix)
        }


def test_excel_translation_replaces_only_text_and_preserves_workbook_objects():
    from app.services.document_translation import translate_document

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "订单"
    sheet["A1"] = "采购订单"
    sheet["A1"].font = Font(name="Microsoft YaHei", size=15, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="5B21B6")
    sheet["A1"].border = Border(bottom=Side(style="thick", color="F59E0B"))
    sheet["B1"] = "数量"
    sheet["B2"] = "=SUM(2,3)"
    sheet["C2"] = 5
    sheet.merge_cells("A3:C3")
    sheet["A3"] = "SC700149043-500"
    sheet.row_dimensions[1].height = 28
    sheet.column_dimensions["A"].width = 24
    image = SpreadsheetImage(BytesIO(_png_bytes()))
    image.width = 48
    image.height = 32
    sheet.add_image(image, "D2")

    source = BytesIO()
    workbook.save(source)
    source_bytes = source.getvalue()

    result = translate_document(
        source_bytes,
        "采购订单.xlsx",
        direction="zh_to_en",
        translator=_fake_translator,
    )

    translated = load_workbook(BytesIO(result.content), data_only=False)
    translated_sheet = translated["订单"]
    assert translated_sheet["A1"].value == "Purchase Order"
    assert translated_sheet["B1"].value == "Quantity"
    assert translated_sheet["B2"].value == "=SUM(2,3)"
    assert translated_sheet["C2"].value == 5
    assert translated_sheet["A3"].value == "SC700149043-500"
    assert translated_sheet["A1"].font.name == "Microsoft YaHei"
    assert translated_sheet["A1"].font.sz == 15
    assert translated_sheet["A1"].font.bold is True
    assert translated_sheet["A1"].fill.fgColor.rgb == "005B21B6"
    assert translated_sheet["A1"].border.bottom.style == "thick"
    assert translated_sheet.merged_cells.ranges == sheet.merged_cells.ranges
    assert translated_sheet.row_dimensions[1].height == 28
    assert translated_sheet.column_dimensions["A"].width == 24
    assert len(translated_sheet._images) == 1
    assert _archive_media(result.content, "xl/media/") == _archive_media(source_bytes, "xl/media/")
    assert result.output_file_name == "采购订单_中译英.xlsx"
    assert result.translated_unit_count == 2


def test_word_translation_preserves_runs_tables_headers_footers_and_images():
    from app.services.document_translation import translate_document

    document = Document()
    paragraph = document.add_paragraph()
    first = paragraph.add_run("采购")
    first.bold = True
    first.font.name = "Microsoft YaHei"
    first.font.size = Pt(12)
    second = paragraph.add_run("订单")
    second.italic = True
    second.font.name = "Arial"
    second.font.size = Pt(16)

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "数量"
    table.cell(0, 1).text = "25"
    section = document.sections[0]
    section.header.paragraphs[0].text = "页眉说明"
    section.footer.paragraphs[0].text = "页脚说明"
    document.add_picture(BytesIO(_png_bytes()), width=Inches(0.6))

    source = BytesIO()
    document.save(source)
    source_bytes = source.getvalue()

    result = translate_document(
        source_bytes,
        "采购订单.docx",
        direction="zh_to_en",
        translator=_fake_translator,
    )

    translated = Document(BytesIO(result.content))
    translated_paragraph = translated.paragraphs[0]
    assert translated_paragraph.text == "Purchase Order"
    assert translated_paragraph.runs[0].text == "Purchase"
    assert translated_paragraph.runs[0].bold is True
    assert translated_paragraph.runs[0].font.name == "Microsoft YaHei"
    assert translated_paragraph.runs[0].font.size == Pt(12)
    assert translated_paragraph.runs[1].text == " Order"
    assert translated_paragraph.runs[1].italic is True
    assert translated_paragraph.runs[1].font.name == "Arial"
    assert translated_paragraph.runs[1].font.size == Pt(16)
    assert translated.tables[0].cell(0, 0).text == "Quantity"
    assert translated.tables[0].cell(0, 1).text == "25"
    assert translated.sections[0].header.paragraphs[0].text == "Header note"
    assert translated.sections[0].footer.paragraphs[0].text == "Footer note"
    assert len(translated.inline_shapes) == 1
    assert _archive_media(result.content, "word/media/") == _archive_media(source_bytes, "word/media/")
    assert result.output_file_name == "采购订单_中译英.docx"


def test_translation_rejects_wrong_direction_content_instead_of_returning_unchanged_file():
    from app.services.document_translation import DocumentTranslationError, translate_document

    workbook = Workbook()
    workbook.active["A1"] = "采购订单"
    source = BytesIO()
    workbook.save(source)

    with pytest.raises(DocumentTranslationError, match="未检测到可翻译的英文文字"):
        translate_document(
            source.getvalue(),
            "订单.xlsx",
            direction="en_to_zh",
            translator=_fake_translator,
        )


def test_excel_translation_supports_direct_string_cells_without_touching_formula_cache():
    from app.services.document_translation import translate_document

    workbook_xml = '''<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets>
</workbook>'''.encode("utf-8")
    workbook_relationships = b'''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Target="worksheets/sheet1.xml"/>
</Relationships>'''
    worksheet_xml = '''<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <sheetData><row r="1">
    <c r="A1" t="str"><v>采购订单</v></c>
    <c r="B1" t="str"><f>TEXT(1,"0")</f><v>数量</v></c>
  </row></sheetData>
</worksheet>'''.encode("utf-8")
    source = BytesIO()
    with ZipFile(source, "w") as archive:
        archive.writestr("xl/workbook.xml", workbook_xml)
        archive.writestr("xl/_rels/workbook.xml.rels", workbook_relationships)
        archive.writestr("xl/worksheets/sheet1.xml", worksheet_xml)

    result = translate_document(
        source.getvalue(),
        "直接字符串.xlsx",
        direction="zh_to_en",
        translator=_fake_translator,
    )

    with ZipFile(BytesIO(result.content)) as archive:
        translated_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "Purchase Order" in translated_xml
    assert "<v>数量</v>" in translated_xml


def test_excel_translation_only_changes_selected_sheets_even_when_shared_strings_are_reused():
    from app.services.document_translation import DocumentTranslationError, translate_document

    workbook_xml = '''<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets>
    <sheet name="华兴报价" sheetId="1" r:id="rId1"/>
    <sheet name="其他厂区" sheetId="2" r:id="rId2"/>
  </sheets>
</workbook>'''.encode("utf-8")
    workbook_relationships = b'''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Target="worksheets/sheet1.xml"/>
  <Relationship Id="rId2" Target="worksheets/sheet2.xml"/>
</Relationships>'''
    shared_strings = '''<?xml version="1.0" encoding="UTF-8"?>
<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="2" uniqueCount="1">
  <si><t>采购订单</t></si>
</sst>'''.encode("utf-8")
    worksheet_xml = '''<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <sheetData><row r="1"><c r="A1" t="s"><v>0</v></c></row></sheetData>
</worksheet>'''.encode("utf-8")
    source = BytesIO()
    with ZipFile(source, "w") as archive:
        archive.writestr("xl/workbook.xml", workbook_xml)
        archive.writestr("xl/_rels/workbook.xml.rels", workbook_relationships)
        archive.writestr("xl/sharedStrings.xml", shared_strings)
        archive.writestr("xl/worksheets/sheet1.xml", worksheet_xml)
        archive.writestr("xl/worksheets/sheet2.xml", worksheet_xml)

    result = translate_document(
        source.getvalue(),
        "多工作表.xlsx",
        direction="zh_to_en",
        translator=_fake_translator,
        selected_sheet_names=["华兴报价"],
    )

    with ZipFile(BytesIO(result.content)) as archive:
        selected_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        untouched_xml = archive.read("xl/worksheets/sheet2.xml")
        output_shared_strings = archive.read("xl/sharedStrings.xml")
    assert 't="inlineStr"' in selected_xml
    assert "Purchase Order" in selected_xml
    assert untouched_xml == worksheet_xml
    assert output_shared_strings == shared_strings
    assert result.translated_unit_count == 1

    with pytest.raises(DocumentTranslationError, match="未找到所选工作表"):
        translate_document(
            source.getvalue(),
            "多工作表.xlsx",
            direction="zh_to_en",
            translator=_fake_translator,
            selected_sheet_names=["不存在"],
        )
