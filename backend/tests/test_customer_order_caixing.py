from __future__ import annotations

from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from fastapi import HTTPException

from app.api.customer_order import _ensure_customer_factory, _validate_upload
from app.services import customer_order_caixing as service
from app.services.customer_order_buzzbee import _decrypt_schedule
from app.services.legacy_excel_bridge import LEGACY_XLS_MAGIC
from starlette.datastructures import UploadFile


PLAYMATES_PO_TEXT = """
Page: 1 of 8
OG-1931815 DATE: 2026/01/23
TO: ROYAL REGENT PRODUCTS (H.K.) LIMITED
CUSTOMER: IMPORTS DRAGON-CAN                                      OUR CONF NO: SG -1926247
PRODUCT NO DESCRIPTION REF C.O. U/M ORDER QTY. UNIT PRICE AMOUNT
HKD HKD
58120E8 WINX CLUB FAIRY WINGS ASST. / PRODUCT LINE: WINX CLUB
Customer Item no.: 58120P
MIX: 58120E8-05
DELIVERY DATE: 2026/04/08 CHN PC
58121E8             WINX CLUB BLOOM FAIRY WINGS ROLE PLAY 300 23.2700 6,981.00
58122E8             WINX CLUB STELLA FAIRY WINGS ROLE PLAY 300 23.2700 6,981.00
ASSORTMENT : 58121E8             2
58122E8             2
8 PCS/CTN
TOTAL : 13,962.00
USE US STANDARD PACKAGING AND US STANDARD INSTRUCTION SHEET WHENEVER APPLICABLE.
GOODS MUST COMPLY WITH ASTM AND EN TOYS SAFETY STANDARD.
"""

SINGLE_PRODUCT_TEXT = """
Page: 1 of 8
OE-1931878 DATE: 2026/02/12
CUSTOMER: PLAYMATES-USA                                        OUR CONF NO: SE -1926304
PRODUCT NO DESCRIPTION REF C.O. U/M ORDER QTY. UNIT PRICE AMOUNT
68695E3 ULTIMATE THUNDER MEGAZORD / PRODUCT LINE: POWER RANGERS
DELIVERY DATE: 2026/05/13 CHN PC 3,000 0.0000 0.00
TOTAL : 0.00
"""

PDQ_PO_TEXT = """
Page: 1 of 8
OG-1932327 DATE: 2026/02/12
CUSTOMER: PLAYMATES-USA                                        OUR CONF NO: SG -1926410
PRODUCT NO DESCRIPTION REF C.O. U/M ORDER QTY. UNIT PRICE AMOUNT
40644E24 COLOR CHANGE MERMAID PDQ / PRODUCT LINE: MERMAID MAGIC
MIX: 40644E24-01
DELIVERY DATE: 2026/05/13 CHN PC
40644A1E24          COLOR CHANGE MERMAID A1 240 10.0000 2,400.00
40644AE24           COLOR CHANGE MERMAID A 240 10.0000 2,400.00
40644B1E24          COLOR CHANGE MERMAID B1 240 10.0000 2,400.00
40644BE24           COLOR CHANGE MERMAID B 240 10.0000 2,400.00
40644C1E24          COLOR CHANGE MERMAID C1 240 10.0000 2,400.00
40644CE24           COLOR CHANGE MERMAID C 240 10.0000 2,400.00
40644D1E24          COLOR CHANGE MERMAID D1 240 10.0000 2,400.00
40644DE24           COLOR CHANGE MERMAID D 240 10.0000 2,400.00
40644E1E24          COLOR CHANGE MERMAID E1 240 10.0000 2,400.00
40644EE24           COLOR CHANGE MERMAID E 240 10.0000 2,400.00
40644FE24           COLOR CHANGE MERMAID F 240 10.0000 2,400.00
ASSORTMENT : 40644A1E24 2
40644AE24 3
40644B1E24 2
40644BE24 2
40644C1E24 2
40644CE24 2
40644D1E24 2
40644DE24 2
40644E1E24 2
40644EE24 3
40644FE24 2
24 PCS/CTN
USE US STANDARD PACKAGING AND US STANDARD INSTRUCTION SHEET WHENEVER APPLICABLE.
"""

RUSSIAN_PACKAGING_TEXT = """
Page: 1 of 9
OG-1932087 DATE: 2026/02/03
CUSTOMER: GULLIVER TOYS                                      OUR CONF NO: SG -1926350
PRODUCT NO DESCRIPTION REF C.O. U/M ORDER QTY. UNIT PRICE AMOUNT
57810GU8 GULLIVER FASHION DOLL ASST.
MIX: 57810GU8-01
DELIVERY DATE: 2026/05/06 CHN PC
57816GU8            GULLIVER FASHION DOLL 240 12.0000 2,880.00
ASSORTMENT : 57816GU8 8
8 PCS/CTN
USE GULLIVER PACKAGING (UNLESS SPECIFY) AND RUS/KZ BI-LINGUAL INSTRUCTION SHEET.
"""


def build_caixing_schedule() -> bytes:
    workbook = openpyxl.Workbook()
    review = workbook.active
    review.title = service.REVIEW_SHEET
    review_headers = {
        "A": "证书",
        "B": "客出单日期",
        "C": "预备单号（OQF NO）",
        "D": "S/C NO",
        "E": "PO.NO",
        "F": "客名/国家",
        "G": "产品编号",
        "H": "产品名称",
        "I": "数量",
        "J": "装箱",
        "K": "箱数",
        "L": "走货数量",
        "M": "剩余未走数量",
        "T": "日期码",
        "U": "客要求走货期",
        "V": "包装要求",
        "W": "国家标准",
        "AA": "单价",
        "AB": "金额HKD",
    }
    for column, value in review_headers.items():
        review[f"{column}2"] = value
    review["G3"] = "58120 E8"
    review["I3"] = "=SUM(I4)"
    review["B4"] = "2026/01/01"
    review["D4"] = "SC-OLD"
    review["E4"] = "PO-OLD"
    review["F4"] = "OLD-CUSTOMER"
    review["G4"] = "58121 E8"
    review["H4"] = "OLD ITEM"
    review["I4"] = 100
    review["J4"] = 2
    review["H5"] = "合计："
    review["I5"] = "=SUM(I4:I4)"
    review["AA5"] = "合计HK$"
    review["AB5"] = "=SUM(AB4:AB4)"
    review["A6"] = service.CURRENT_ORDER_MARKER

    order = workbook.create_sheet(service.ORDER_SHEET)
    order_headers = {
        "A": "FA",
        "B": "客出单日期",
        "C": "预备单号（OQF NO）",
        "D": "S/C NO",
        "E": "PO.NO",
        "F": "客名/国家",
        "G": "产品编号",
        "H": "产品名称",
        "I": "数量",
        "J": "装箱",
        "K": "单价",
        "L": "金额",
        "M": "客要求走货期",
        "AD": 58121,
        "AE": 58122,
    }
    for column, value in order_headers.items():
        order[f"{column}2"] = value
    order["G3"] = "58120 E8"
    order["I3"] = 100
    order["B4"] = "2026/01/01"
    order["D4"] = "SC-OLD"
    order["E4"] = "PO-OLD"
    order["G4"] = "58121 E8"
    order["I4"] = 100
    order["A5"] = service.CURRENT_ORDER_MARKER
    order["M5"] = "合计："
    order["AD5"] = "=SUM(AD3:AD4)"
    order["AE5"] = "=SUM(AE3:AE4)"

    item = workbook.create_sheet(service.ITEM_SHEET)
    item["AG1"] = 58121
    item["AH1"] = 58122
    item_headers = {
        "A": "证书",
        "B": "客出单日期",
        "C": "预备单号（OQF NO）",
        "D": "S/C NO",
        "E": "PO.NO",
        "F": "客名/国家",
        "G": "产品编号",
        "H": "产品名称",
        "I": "数量",
        "J": "包装",
        "K": "备注",
        "O": "客要求走货期",
    }
    for column, value in item_headers.items():
        item[f"{column}3"] = value
    item["G4"] = "58120 E8"
    item["I4"] = 100
    item["B5"] = "2026/01/01"
    item["D5"] = "SC-OLD"
    item["E5"] = "PO-OLD"
    item["G5"] = "58121 E8"
    item["I5"] = 100
    item["A6"] = service.CURRENT_ORDER_MARKER
    item["P6"] = "合计："
    item["AG6"] = "=SUM(AG4:AG5)"
    item["AH6"] = "=SUM(AH4:AH5)"

    review.freeze_panes = "J3"
    order.freeze_panes = "A3"
    item.freeze_panes = "J4"

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_caixing_parser_extracts_groups_delivery_pack_and_adjusted_price(monkeypatch):
    monkeypatch.setattr(service, "_extract_pdf_text", lambda _content: PLAYMATES_PO_TEXT)

    rows = service.parse_caixing_pdf("1931815.pdf", b"%PDF synthetic")

    assert len(rows) == 2
    assert rows[0].source_po_date == "2026/01/23"
    assert rows[0].requested_ship_date == "2026/04/08"
    assert rows[0].po_no == "OG-1931815"
    assert rows[0].contract_no == "SG-1926247"
    assert rows[0].customer_name == "IMPORTS DRAGON-CAN"
    assert rows[0].parent_product_no == "58120 E8"
    assert rows[0].parent_quantity == 600
    assert rows[0].parent_units_per_carton == 8
    assert rows[0].product_no == "58121 E8"
    assert rows[0].quantity == 300
    assert rows[0].units_per_carton == 2
    assert rows[0].raw_unit_price_hkd == Decimal("23.2700")
    assert rows[0].unit_price_hkd == Decimal("22.2228500")
    assert rows[0].amount_hkd == Decimal("6666.8550000")
    assert rows[0].packaging == "美版彩盒"
    assert rows[0].standard == "美国标准"
    assert rows[1].product_no == "58122 E8"
    assert rows[1].units_per_carton == 2


def test_caixing_parser_handles_quantity_on_delivery_date_line(monkeypatch):
    monkeypatch.setattr(service, "_extract_pdf_text", lambda _content: SINGLE_PRODUCT_TEXT)

    rows = service.parse_caixing_pdf("1931878.pdf", b"%PDF synthetic")

    assert len(rows) == 1
    assert rows[0].product_no == "68695 E3"
    assert rows[0].parent_product_no == "68695 E3"
    assert rows[0].quantity == 3000
    assert rows[0].units_per_carton == 3
    assert rows[0].parent_units_per_carton == 3
    assert rows[0].requested_ship_date == "2026/05/13"
    assert rows[0].raw_unit_price_hkd == 0
    assert rows[0].amount_hkd == 0


def test_caixing_parser_keeps_full_pdq_assortment_product_codes(monkeypatch):
    monkeypatch.setattr(service, "_extract_pdf_text", lambda _content: PDQ_PO_TEXT)

    rows = service.parse_caixing_pdf("1932327.pdf", b"%PDF synthetic")
    packs = {row.product_no: row.units_per_carton for row in rows}

    assert len(rows) == 11
    assert {row.parent_product_no for row in rows} == {"40644 E24"}
    assert packs["40644 AE24"] == 3
    assert packs["40644 EE24"] == 3
    assert packs["40644 A1E24"] == 2
    assert packs["40644 FE24"] == 2
    assert sum(packs.values(), Decimal("0")) == 24


def test_caixing_parser_recognizes_gulliver_russian_packaging(monkeypatch):
    monkeypatch.setattr(
        service,
        "_extract_pdf_text",
        lambda _content: RUSSIAN_PACKAGING_TEXT,
    )

    rows = service.parse_caixing_pdf("1932087.pdf", b"%PDF synthetic")

    assert len(rows) == 1
    assert rows[0].parent_product_no == "57810 GU8"
    assert rows[0].packaging == "俄罗斯彩盒包装"
    assert rows[0].standard == ""


def test_caixing_preview_and_export_write_review_order_and_item(monkeypatch):
    monkeypatch.setattr(service, "_extract_pdf_text", lambda _content: PLAYMATES_PO_TEXT)
    schedule_content = build_caixing_schedule()
    po_files = [("1931815.pdf", b"%PDF synthetic")]

    preview = service.create_caixing_batch_preview(
        factory_id="huaxing",
        received_date="2026-07-29",
        po_files=po_files,
        schedule_file_name="2026年彩星排期.xlsx",
        schedule_content=schedule_content,
    )

    assert preview["customer_code"] == "caixing"
    assert preview["summary"] == {
        "total": 3,
        "valid": 3,
        "warning": 0,
        "blocked": 0,
    }
    assert preview["preview_schema_version"] == "customer-order-caixing-preview-v2"
    assert preview["target_template"] == service.TARGET_TEMPLATE
    assert preview["output_file_name"] == "2026年彩星排期.xlsx"
    assert preview["rows"][0]["row_role"] == "parent"
    assert preview["rows"][0]["product_no"] == "58120 E8"
    assert preview["rows"][0]["quantity"] == "600"
    assert preview["rows"][0]["units_per_carton"] == "8"
    assert preview["rows"][0]["carton_count"] == "75"
    assert preview["rows"][1]["row_role"] == "detail"
    assert preview["rows"][1]["parent_product_no"] == "58120 E8"
    assert preview["rows"][1]["received_date"] == "2026-07-29"
    assert preview["rows"][1]["requested_ship_date"] == "2026-04-08"
    assert preview["rows"][1]["units_per_carton"] == "2"
    assert preview["rows"][1]["carton_count"] == "150"
    assert preview["rows"][1]["item_sheet_name"] == "正单评审表 / 接单表 / ITEM表"

    output, file_name, _ = service.export_caixing_batch_schedule(
        factory_id="huaxing",
        received_date="2026-07-29",
        po_files=po_files,
        schedule_file_name="2026年彩星排期.xlsx",
        schedule_content=schedule_content,
    )

    assert file_name == "2026年彩星排期.xlsx"
    plain, encrypted = _decrypt_schedule(output)
    assert encrypted is False
    rendered = openpyxl.load_workbook(BytesIO(plain), data_only=False)
    assert rendered[service.REVIEW_SHEET]["G5"].fill.fgColor.rgb == "FFFFFF00"
    assert rendered[service.ORDER_SHEET]["G5"].fill.fgColor.rgb == "FFFFFF00"
    assert rendered[service.ITEM_SHEET]["G6"].fill.fgColor.rgb == "FFFFFF00"
    assert rendered[service.REVIEW_SHEET]["G4"].fill.fgColor.rgb != "FFFFFF00"
    assert rendered[service.REVIEW_SHEET].freeze_panes == "J3"
    assert rendered[service.ORDER_SHEET].freeze_panes == "A3"
    assert rendered[service.ITEM_SHEET].freeze_panes == "J4"
    rendered.close()
    workbook = service.CaixingSchedule(plain)

    review_rows = workbook.read_rows(service.REVIEW_SHEET)
    assert review_rows[5]["G"] == "58120 E8"
    assert review_rows[5]["H"] == "WINX CLUB FAIRY WINGS ASST. / PRODUCT LINE: WINX CLUB"
    assert review_rows[5]["I"] == "600"
    assert review_rows[5]["J"] == "8"
    assert workbook.read_cell_formula(service.REVIEW_SHEET, 5, "I") == "SUM(I6:I7)"
    assert review_rows[6]["B"] == "2026/07/29"
    assert review_rows[6]["E"] == "OG-1931815"
    assert review_rows[6]["J"] == "2"
    assert review_rows[6]["K"] == "150"
    assert review_rows[6]["M"] == "300"
    assert review_rows[6]["U"] == "2026/04/08"
    assert review_rows[6]["V"] == "美版彩盒"
    assert review_rows[6]["W"] == "美国标准"
    assert review_rows[6].get("T", "") == ""
    assert review_rows[6]["AA"] == "22.22285"
    assert review_rows[6]["AB"] == "6666.855"

    order_rows = workbook.read_rows(service.ORDER_SHEET)
    assert order_rows[5]["G"] == "58120 E8"
    assert order_rows[5]["H"] == "WINX CLUB FAIRY WINGS ASST. / PRODUCT LINE: WINX CLUB"
    assert order_rows[5]["I"] == "600"
    assert order_rows[5]["J"] == "8"
    assert order_rows[6]["E"] == "OG-1931815"
    assert order_rows[6]["AD"] == "300"
    assert order_rows[7]["AE"] == "300"
    assert order_rows[6]["M"] == "2026/04/08"

    item_rows = workbook.read_rows(service.ITEM_SHEET)
    assert item_rows[6]["G"] == "58120 E8"
    assert item_rows[6]["H"] == "WINX CLUB FAIRY WINGS ASST. / PRODUCT LINE: WINX CLUB"
    assert item_rows[6]["I"] == "600"
    assert item_rows[7]["E"] == "OG-1931815"
    assert item_rows[7]["J"] == "美版彩盒"
    assert item_rows[7]["AG"] == "300"
    assert item_rows[8]["AH"] == "300"
    assert item_rows[7]["O"] == "2026/04/08"


def test_caixing_three_po_batch_allows_missing_matrix_columns(monkeypatch):
    texts = {
        b"standard": PLAYMATES_PO_TEXT,
        b"pdq": PDQ_PO_TEXT,
        b"single": SINGLE_PRODUCT_TEXT,
    }
    monkeypatch.setattr(service, "_extract_pdf_text", lambda content: texts[content])
    schedule_content = build_caixing_schedule()
    po_files = [
        ("1932232.pdf", b"standard"),
        ("1932327.pdf", b"pdq"),
        ("1932524.pdf", b"single"),
    ]

    preview = service.create_caixing_batch_preview(
        factory_id="huaxing",
        received_date="2026-07-29",
        po_files=po_files,
        schedule_file_name="2026年彩星排期.xlsx",
        schedule_content=schedule_content,
    )

    matrix_issues = [
        issue
        for row in preview["rows"]
        for issue in row["issues"]
        if issue["code"] == "missing_product_matrix_column"
    ]
    assert preview["summary"]["blocked"] == 0
    assert matrix_issues
    assert all(issue["severity"] == "warning" for issue in matrix_issues)
    assert all(issue["can_skip"] is False for issue in matrix_issues)

    output, file_name, exported_preview = service.export_caixing_batch_schedule(
        factory_id="huaxing",
        received_date="2026-07-29",
        po_files=po_files,
        schedule_file_name="2026年彩星排期.xlsx",
        schedule_content=schedule_content,
    )

    assert output.startswith(b"PK")
    assert file_name == "2026年彩星排期.xlsx"
    assert exported_preview["summary"]["blocked"] == 0


def test_caixing_existing_order_line_warns_but_allows_test_export(monkeypatch):
    monkeypatch.setattr(service, "_extract_pdf_text", lambda _content: PLAYMATES_PO_TEXT)
    source = openpyxl.load_workbook(BytesIO(build_caixing_schedule()))
    order = source[service.ORDER_SHEET]
    order["D4"] = "SG-1926247"
    order["E4"] = "OG-1931815"
    order["G4"] = "58121 E8"
    output = BytesIO()
    source.save(output)
    source.close()
    schedule_content = output.getvalue()
    po_files = [("1931815.pdf", b"%PDF synthetic")]

    preview = service.create_caixing_batch_preview(
        factory_id="huaxing",
        received_date="2026-07-29",
        po_files=po_files,
        schedule_file_name="2026年彩星排期.xlsx",
        schedule_content=schedule_content,
    )

    duplicate_row = next(
        row for row in preview["rows"] if row["product_no"] == "58121 E8"
    )
    duplicate_issue = next(
        issue
        for issue in duplicate_row["issues"]
        if issue["code"] == "existing_order_line"
    )
    assert duplicate_row["status"] == "warning"
    assert duplicate_issue["severity"] == "warning"
    assert duplicate_issue["can_skip"] is False
    assert preview["summary"] == {
        "total": 3,
        "valid": 2,
        "warning": 1,
        "blocked": 0,
    }

    exported, file_name, exported_preview = service.export_caixing_batch_schedule(
        factory_id="huaxing",
        received_date="2026-07-29",
        po_files=po_files,
        schedule_file_name="2026年彩星排期.xlsx",
        schedule_content=schedule_content,
    )
    assert exported.startswith(b"PK")
    assert file_name == "2026年彩星排期.xlsx"
    assert exported_preview["summary"]["blocked"] == 0


def test_caixing_legacy_xls_uses_converter_and_preserves_format(monkeypatch):
    monkeypatch.setattr(service, "_extract_pdf_text", lambda _content: PLAYMATES_PO_TEXT)
    legacy_schedule = LEGACY_XLS_MAGIC + b"encrypted-schedule"
    converted_schedule = build_caixing_schedule()
    converted_outputs: list[tuple[bytes, str | None]] = []

    monkeypatch.setattr(service, "_decrypt_schedule", lambda content: (legacy_schedule, True))
    monkeypatch.setattr(
        service,
        "convert_legacy_xls_to_xlsx",
        lambda content: converted_schedule,
    )

    def fake_convert_to_xls(content: bytes, *, output_password: str | None) -> bytes:
        converted_outputs.append((content, output_password))
        return LEGACY_XLS_MAGIC + b"converted-output"

    monkeypatch.setattr(service, "convert_xlsx_to_legacy_xls", fake_convert_to_xls)
    po_files = [("1931815.pdf", b"%PDF synthetic")]

    preview = service.create_caixing_batch_preview(
        factory_id="huaxing",
        received_date="2026-07-29",
        po_files=po_files,
        schedule_file_name="2026年彩星生产排期表.xls",
        schedule_content=b"encrypted-upload",
    )

    assert preview["output_file_name"] == "2026年彩星生产排期表.xls"
    assert preview["_schedule_encrypted"] is True
    assert preview["_schedule_format"] == "xls"
    assert any("Excel 97-2003 .xls" in warning for warning in preview["warnings"])

    output, file_name, _ = service.export_caixing_batch_schedule(
        factory_id="huaxing",
        received_date="2026-07-29",
        po_files=po_files,
        schedule_file_name="2026年彩星生产排期表.xls",
        schedule_content=b"encrypted-upload",
    )

    assert output == LEGACY_XLS_MAGIC + b"converted-output"
    assert file_name == "2026年彩星生产排期表.xls"
    assert len(converted_outputs) == 1
    assert converted_outputs[0][0].startswith(b"PK")
    assert converted_outputs[0][1] == "2026"


def test_caixing_api_accepts_xls_schedule_extension():
    _validate_upload(
        UploadFile(
            filename="2026年彩星生产排期表.xls",
            file=BytesIO(b"schedule"),
        ),
        kind="客户排期",
        supported=(".xls", ".xlsx"),
    )


def test_caixing_customer_is_owned_by_huaxing():
    _ensure_customer_factory("caixing", "huaxing")
    with pytest.raises(HTTPException, match="彩星.*只属于华兴厂区"):
        _ensure_customer_factory("caixing", "huakang-a")
