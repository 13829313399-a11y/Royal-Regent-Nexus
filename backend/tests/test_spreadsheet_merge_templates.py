"""Real parser contracts for six public-tool template layouts and isolated audit data."""
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
import pytest

from app.schemas.spreadsheet_merge import MergeRequest
from app.services import spreadsheet_merge as merge
from app.services import spreadsheet_merge_templates as templates

PUBLIC = Path(__file__).resolve().parents[2] / 'public' / 'templates'
STATIC = {
    'carton-history-orders': 'carton-history-order-import-template.xlsx',
    'carton-opening-inventory': 'carton-history-inventory-import-template.xlsx',
    'carton-weekly-schedule': 'carton-weekly-schedule-template.xlsx',
}


def payload(template_id, rows, **extra):
    info = templates.get_template(template_id)
    return MergeRequest.model_validate({
        'template_id': template_id, 'factory_id': 'huaxing', 'profile': info['profile'],
        'primary': {'name': info['label'], 'headers': info['headers'], 'header_row': info['header_row'], 'rows': []},
        'secondary': {'name': '原文件.xlsx/原始数据', 'headers': info['headers'], 'rows': rows},
        'mapping': [{'source': index, 'target': index} for index in range(len(info['headers']))], **extra,
    })


def examples(template_id, number):
    book = load_workbook(BytesIO(templates.download_template(template_id)))
    info = templates.get_template(template_id)
    row = list(book['填写示例'].values)[number - 1]
    return [merge._scalar(value) for value in row[:len(info['headers'])]]


def test_catalogue_is_bounded_native_and_has_no_business_queries(monkeypatch):
    from sqlalchemy.orm import Session
    def forbidden(*args, **kwargs):
        raise AssertionError('Template operation must not query any business database')
    monkeypatch.setattr(Session, 'execute', forbidden)
    monkeypatch.setattr(Session, 'scalars', forbidden)
    records = templates.list_templates()
    assert [record['id'] for record in records] == [
        'carton-packaging', 'carton-paper-options', 'carton-locations',
        'carton-history-orders', 'carton-opening-inventory', 'carton-weekly-schedule',
    ]
    assert [record['header_row'] for record in records] == [1, 1, 1, 2, 2, 3]
    assert [record['profile'] for record in records] == ['carton-packaging'] + ['generic'] * 5
    for record in records:
        assert record['description'] and record['headers'] and record['sheet_name']
        book = load_workbook(BytesIO(templates.download_template(record['id'])))
        main = book[record['sheet_name']]
        assert [c.value for c in main[record['header_row']]] == record['headers']
        assert not any(c.value is not None for row in main.iter_rows(min_row=record['header_row'] + 1) for c in row)
    changed = templates.get_template('carton-locations')
    changed['headers'][0] = '不能改变全局'
    assert templates.get_template('carton-locations')['headers'][0] == '仓库'


@pytest.mark.parametrize('template_id,filename', STATIC.items())
def test_packaged_assets_are_exact_canonical_bytes_and_need_no_frontend_directory(template_id, filename, monkeypatch):
    canonical = (PUBLIC / filename).read_bytes()
    assert templates.download_template(template_id) == canonical
    assert templates.BUNDLED_TEMPLATES.is_relative_to(Path(__file__).resolve().parents[1] / 'app')
    original = Path.read_bytes
    def backend_only(path):
        if path.is_relative_to(PUBLIC):
            raise AssertionError('The API container has no public directory')
        return original(path)
    monkeypatch.setattr(Path, 'read_bytes', backend_only)
    assert templates.download_template(template_id) == canonical


def test_registry_rejects_unknown_paths_and_selection_contract_changes():
    for identifier in ('../carton-history-order-import-template.xlsx', 'locations', '/etc/passwd'):
        with pytest.raises(HTTPException) as error:
            templates.download_template(identifier)
        assert error.value.status_code == 404
    request = payload('carton-history-orders', [])
    with pytest.raises(HTTPException):
        merge.merge(request.model_copy(update={'profile': 'carton-packaging'}))
    request.primary.headers[0] = '错误列名'
    with pytest.raises(HTTPException):
        merge.merge(request)


@pytest.mark.parametrize('template_id,kind,row', [
    ('carton-paper-options', 'paper-options', ['外箱', 'A33', 30, 20, 10]),
    ('carton-locations', 'locations', ['A仓', 'B01-B03']),
    ('carton-packaging', 'configurations', ['00123', '产品', '组', '外箱', 'A33', 30, 20, 10, 'cm', '个', 24, '否', '启用', '', None, None]),
])
def test_master_exports_parse_only_real_input_and_keep_canonical_examples(template_id, kind, row):
    from app.services.carton_master_import import parse
    request = payload(template_id, [row])
    result = merge.merge(request)
    assert result['ready_for_import'] == (template_id == 'carton-packaging')
    content, filename, _ = merge.export(request)
    assert filename.endswith('ready.xlsx' if template_id == 'carton-packaging' else 'pending.xlsx')
    book = load_workbook(BytesIO(content))
    assert book.sheetnames == ['导入数据', '说明与示例']
    entries, errors = parse(content, 'export.xlsx', kind)
    assert len(entries) == 1 and errors == []
    assert book['说明与示例'].max_row > 20


def test_history_wide_layout_title_supplement_and_examples_are_preserved_but_audit_never_imports():
    from app.services.carton_procurement_history_import import _parse_rows
    request = payload('carton-history-orders', [examples('carton-history-orders', 6)])
    assert not merge.merge(request)['ready_for_import']
    content, filename, _ = merge.export(request)
    assert filename.endswith('pending.xlsx')
    book = load_workbook(BytesIO(content))
    assert book.sheetnames == ['历史订单导入', '附加纸品明细', '填写示例', '字段说明']
    assert book['历史订单导入']['A1'].value == '历史订单导入'
    assert book['历史订单导入']['A2'].value == '客户名称'
    assert not any(c.value is not None for row in book['附加纸品明细'].iter_rows(min_row=3) for c in row)
    parsed, _ = _parse_rows('history.xlsx', content)
    assert len(parsed) == 1 and parsed[0]['item_no'] == '000203307004'
    assert parsed[0]['product_name'] == '' and parsed[0]['product_order_quantity'] is None
    assert '待补版' in str(list(book['字段说明'].values))


def test_opening_layout_uses_per_row_time_and_never_imports_examples_or_audit():
    from app.services.carton_opening_preview import options_from_json, parse_rows
    request = payload('carton-opening-inventory', [examples('carton-opening-inventory', 7)])
    content, filename, _ = merge.export(request)
    assert filename.endswith('pending.xlsx')
    book = load_workbook(BytesIO(content))
    assert book.sheetnames == ['期初库存导入', '填写示例', '字段说明']
    assert book['期初库存导入']['O2'].value == '入库时间'
    parsed, _ = parse_rows('opening.xlsx', content, options_from_json({'dimension_unit': 'cm', 'warehouse': 'A', 'customer_name': '迪奇'}))
    assert len(parsed) == 1 and parsed[0]['unit_price'] is None and parsed[0]['amount'] is None
    assert parsed[0]['original_inbound_at'].startswith('2026-')


def test_weekly_layout_and_compact_full_audit_do_not_create_phantom_business_rows():
    from app.services.carton_procurement_imports import _parse_weekly
    row = ['0010001234/200', '0001', 'DISPLAY-ONLY', '迪奇', '000203307004', '', 3600, '1/120', '2026-09-21']
    # Direct source headers would match both Reference and 数量 if exposed as an audit row.
    # More than 5000 original audit records also exceed the business parser's physical row limit.
    audit_rows = [row for _ in range(5000)]
    request = payload('carton-weekly-schedule', [row], audit_tables=[{
        'name': '完整原文索引', 'headers': templates.get_template('carton-weekly-schedule')['headers'], 'rows': audit_rows,
    }])
    content, filename, _ = merge.export(request)
    assert filename.endswith('pending.xlsx')
    book = load_workbook(BytesIO(content))
    assert book.sheetnames == ['导入数据', '填写说明与示例']
    assert book['导入数据']['A3'].value == 'Reference'
    assert book['导入数据']['A4'].value == row[0]
    assert book['导入数据']['A2'].value.startswith('请从第 4 行')
    parsed = _parse_weekly('schedule.xlsx', content)
    assert len(parsed['rows']) == 1 and not parsed['warnings']
    assert parsed['rows'][0]['customer_po'] == '0001'
    assert parsed['rows'][0]['po_numbers'] == 'DISPLAY-ONLY'
    assert parsed['rows'][0]['source_row'] == 4
    notes = book['填写说明与示例']
    marker = next(cell.row for row_cells in notes for cell in row_cells if isinstance(cell.value, str) and cell.value.startswith('转换审计：'))
    stream = ''.join(str(notes.cell(r, 1).value or '') for r in range(marker + 1, notes.max_row + 1))
    records = [json.loads(line) for line in stream.splitlines()]
    assert sum(record[0] == '审计附表数据' for record in records) == 5000
    assert next(record for record in records if record[0] == '输出行来源')[1] == 4
    assert notes.max_row < 5001 and notes.max_column < 256


def test_selected_template_native_row_limits_require_explicit_batches():
    request = payload('carton-locations', [['A仓', f'B{i:04}'] for i in range(1001)])
    with pytest.raises(HTTPException, match='每批|分批'):
        merge.export(request)


def test_batched_template_audit_references_actual_native_instruction_sheet():
    first, second = examples('carton-history-orders', 6), examples('carton-history-orders', 7)
    content, _, _ = merge.export(payload('carton-history-orders', [first, second], batch_size=1))
    with ZipFile(BytesIO(content)) as archive:
        book = load_workbook(BytesIO(archive.read(archive.namelist()[1])))
    assert '第 001 批文件的字段说明页' in str(list(book['字段说明'].values))
    notes = book['字段说明']
    marker = next(cell.row for row_cells in notes for cell in row_cells if isinstance(cell.value, str) and cell.value.startswith('转换审计：'))
    records = [json.loads(line) for line in ''.join(str(notes.cell(r, 1).value or '') for r in range(marker + 1, notes.max_row + 1)).splitlines()]
    reference = next(record for record in records if record[0] == '输出行来源')
    assert reference[1] == 3 and reference[4] == 3 and reference[6] == 3
