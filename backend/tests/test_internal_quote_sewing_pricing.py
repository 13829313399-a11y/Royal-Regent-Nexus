import json
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.services.internal_quote_calculator import CalculationInputError
from app.services.internal_quote_excel import _build_sewing_sheet
from app.services.internal_quote_import import parse_internal_quote_workbook
from app.services.internal_quote_sewing import sewing_formula_multiplier, sewing_workbook_formulas
from app.services.internal_quote_templates import build_internal_quote_import_template
from app.services.internal_quote_workbook_view import WorkbookValues
from test_internal_quote_sewing_currency import SNAPSHOT, calculate, parse_rows


HEADER = ['物料名称', '部位', '用量', '单价', '汇率', '成本', '码点', '价钱']


def test_pending_price_beats_cost_column_and_only_recalculates_when_both_inputs_exist():
    payload = parse_rows(HEADER, [['IC灯一套', '', None, None, .85, 0, 1.1, 17.55]])
    row = payload['groups'][0]['materials'][0]
    assert row['usage'] is None and row['unit_price_rmb'] is None
    assert row['quoted_price_hkd'] == '17.5500'
    row['extra_multiplier'] = '2'
    result = calculate(payload)
    assert result['status'] == 'valid'
    assert result['totals']['total_hkd'] == '17.5500'
    assert result['warnings'][0]['severity'] == 'warning'
    row['unit_price_rmb'] = '8.5'
    assert calculate(payload)['totals']['total_hkd'] == '17.5500'
    row['usage'] = 1
    result = calculate(payload)
    assert result['totals']['total_hkd'] == '22.0000'
    assert not result['warnings']
    row['unit_price_rmb'] = 0  # An explicit zero is not a missing input.
    assert calculate(payload)['totals']['total_hkd'] == '0.0000'


def test_formula_multipliers_are_per_row_editable_and_labor_is_not_material():
    payload = parse_rows(HEADER, [
        ['裁床', '', 1, .85, .85, None, 1.1, '=F2*G2*1.72'],
        ['车缝', '', 1, 1.7, .85, None, 1.1, '=$F$3*$G$3*1.5'],
        ['手工', '', 1, .85, .85, None, 1, '=F4*G4*1.25'],
    ])
    rows = payload['groups'][0]['materials']
    assert [r['extra_multiplier'] for r in rows] == ['1.7200', '1.5000', '1.2500']
    result = calculate(payload)
    assert result['totals']['total_hkd'] == '6.4420'
    assert result['totals']['labor_hkd'] == '6.4420'
    rows[1]['extra_multiplier'] = '1.8'
    assert calculate(payload)['totals']['total_hkd'] == '7.1020'
    for invalid in ('-1', '0', 'NaN', 'Infinity'):
        rows[1]['extra_multiplier'] = invalid
        with pytest.raises(CalculationInputError):
            calculate(payload)


def test_multiplier_reader_never_evaluates_other_cells_functions_or_external_references():
    for formula in ('=F3*G2*1.5', '=F2*G2*H2', '=F2*G2*WEBSERVICE("https://invalid")', '=other!F2*G2*1.5'):
        assert sewing_formula_multiplier(formula, 2, 5, 6) is None
    content, _ = build_internal_quote_import_template('hair')
    # The checked-in native BIFF8 template exercises the XLS token reader;
    # private customer workbooks must never become repository fixtures.
    formulas = sewing_workbook_formulas(content, '明细', 100, 30)
    assert formulas
    assert all(isinstance(key, tuple) for key in formulas)


def test_export_pending_price_and_multiplier_roundtrip_without_formula_caches():
    payload = parse_rows(HEADER, [
        ['IC灯', '', None, None, .85, 0, 1.1, 17.55],
        ['车缝', '', 1, 1.7, .85, None, 1.1, '=F3*G3*1.5'],
    ])
    book = Workbook()
    _build_sewing_sheet(book, SimpleNamespace(payload_json=json.dumps(payload)), SNAPSHOT)
    sheet = book['车缝明细']
    assert sheet['F5'].value == '' and sheet['G5'].value == ''
    assert sheet['O5'].value == 17.55 and sheet['N6'].value == 1.5
    values = WorkbookValues(book)
    assert float(values.cell('车缝明细', 'K5')) == pytest.approx(17.55)
    assert float(values.cell('车缝明细', 'K6')) == pytest.approx(3.3)
    sheet['F5'] = 1
    sheet['G5'] = 0
    assert WorkbookValues(book).cell('车缝明细', 'K5') == 0
    sheet['F5'] = ''
    sheet['G5'] = ''
    output = BytesIO()
    book.save(output)
    imported = parse_internal_quote_workbook(output.getvalue(), 'sewing').payload_fragment
    assert calculate(imported)['totals'] == calculate(payload)['totals']


def test_import_confirm_edit_and_reload_retain_pending_price_and_variable_multiplier(monkeypatch):
    from test_internal_quote_api import ALL_SECTION_CODES, create_payload, login, make_client

    with make_client(monkeypatch) as client:
        login(client, 'sewing-price-owner', 'admin', '*', '*')
        created = client.post('/api/internal-quotes', json=create_payload(
            suffix='SEWING-PRICING', participating_sections=ALL_SECTION_CODES,
        ))
        assert created.status_code == 201, created.text
        quote_id = created.json()['id']
        book = Workbook()
        book.active.append(HEADER)
        book.active.append(['IC灯', '', None, None, .85, 0, 1.1, 17.55])
        book.active.append(['车缝', '', 1, 1.7, .85, None, 1.1, '=F3*G3*1.5'])
        output = BytesIO()
        book.save(output)
        preview = client.post(f'/api/internal-quotes/{quote_id}/imports/sewing/preview', files={
            'file': ('sewing.xlsx', output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
        })
        assert preview.status_code == 201, preview.text
        batch = preview.json()
        response = client.post(f"/api/internal-quotes/{quote_id}/imports/{batch['batch_id']}/confirm", json={
            'revision': batch['target_revision'], 'mode': 'replace',
        })
        assert response.status_code == 200, response.text
        sewing = next(s for s in client.get(f'/api/internal-quotes/{quote_id}').json()['sections'] if s['department'] == 'sewing')
        assert sewing['calculation']['totals']['total_hkd'] == '20.8500'
        rows = sewing['payload']['groups'][0]['materials']
        assert rows[0]['unit_price_rmb'] is None
        rows[1]['extra_multiplier'] = '1.8'
        response = client.put(f'/api/internal-quotes/{quote_id}/sections/sewing', json={
            'revision': sewing['revision'], 'payload': sewing['payload'],
        })
        assert response.status_code == 200, response.text
        sewing = next(s for s in client.get(f'/api/internal-quotes/{quote_id}').json()['sections'] if s['department'] == 'sewing')
        assert sewing['calculation']['totals']['total_hkd'] == '21.5100'
        assert sewing['payload']['groups'][0]['materials'][1]['extra_multiplier'] == '1.8'
        layout = client.get(f'/api/internal-quotes/{quote_id}/output-layout')
        assert layout.status_code == 200, layout.text
        assert any(w.get('code') == 'sewing_quoted_price_pending' for w in layout.json()['warnings'])
        group = next(g for g in layout.json()['groups'] if g['label'] == 'IC灯')
        price_field = next(f for f in group['fields'] if f['kind'] == 'price')
        response = client.put(f'/api/internal-quotes/{quote_id}/output-layout', json={
            'base_token': layout.json()['base_token'], 'edits': [{'field_id': price_field['id'], 'value': '8.5'}],
        })
        assert response.status_code == 200, response.text
        sewing = next(s for s in client.get(f'/api/internal-quotes/{quote_id}').json()['sections'] if s['department'] == 'sewing')
        assert sewing['payload']['groups'][0]['materials'][0]['usage'] is None
        assert sewing['calculation']['totals']['total_hkd'] == '21.5100'


def test_source_price_field_preserves_hkd_currency():
    from app.services.internal_quote_output_fields import source_groups
    section = SimpleNamespace(is_required=True, department='sewing', department_name='车缝', payload_json=json.dumps({
        'groups': [{'materials': [{'item': '布', 'unit_price_source_currency': 'HKD', 'unit_price_hkd': 3.5}]}],
    }))
    fields = source_groups([section])[0]['fields']
    assert next(f for f in fields if f['kind'] == 'price')['path'][-1] == 'unit_price_hkd'
