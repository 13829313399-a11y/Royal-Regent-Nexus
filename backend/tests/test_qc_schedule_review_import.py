from datetime import datetime
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.utils.datetime import to_excel

from test_molding_sample_api import login_as, make_client


HEADERS = [
    "客出单日期", "订单类型", "Contract No.", "SO#/Reference", "P/O#:", "客名",
    "产品编号", "产品中文名称", "产品英文名称", "數量", "装箱", "箱数",
    "国家标准", "单价HK", "金额HK", "单价USD", "总金额USD", "验货日期",
    "客要求走货期", "备注",
]


def review_workbook(*, boundary=True, review_sheet=True, order_type=True):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "正单评审表" if review_sheet else "接单表"
    sheet.append(["正单评审表"])
    sheet.append(["操作说明"])
    headers = HEADERS.copy()
    if not order_type:
        headers[1] = "其他"
    sheet.append(headers)
    sheet.append([
        datetime(2026, 7, 9), "正式PO", "053128", "INTERNAL-SO", 6661356, "WMC",
        "0045841", "小链条枪", "BATTLE BLAZER", "3,752", "0/4", 938,
        "加拿大标准", None, None, None, None, datetime(2026, 9, 15),
        to_excel(datetime(2026, 9, 21)), "取消单",  # A remark is not the section boundary.
    ])
    sheet["E4"].number_format = "0000000000"
    sheet.append([None, "备料单", "IGNORE-STOCK", None, "IGNORE-PO", "WMC", "001", None, None, 9])
    sheet.append([None, "正单", "053129", None, "000123", "WMU", "0045842", "产品二", None, 120,
                  "6P/12", 10, None, None, None, None, None, "2026/09/17", "2026-09-30"])
    sheet.append([])
    sheet.append([None, "正式PO", "053130", None, None, "WMU", "0045843", "缺PO保留", None, 1,
                  None, None, None, None, None, None, None, "9/21-10% 10/23-80%"])
    if boundary:
        sheet.append([" 取消单 "])
        sheet.merge_cells("A9:T9")
        sheet.append([None, "正式PO", "CANCELLED", None, "999", "WMC", "999", None, None, 999])
        sheet.append(["已走货订单"])
        sheet.append([None, "正式PO", "SHIPPED", None, "999", "WMC", "999", None, None, 999])
    other = workbook.create_sheet("ITEM表-套装")
    other.append(HEADERS)
    other.append([None, "正式PO", "ITEM-DUPLICATE", None, "999", "WMC", "999", None, None, 999])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def parse(content):
    from app.services.qc_inspection import _parse_schedule_rows
    return _parse_schedule_rows(content)


def test_only_review_formal_orders_above_cancelled_with_source_values():
    rows = parse(review_workbook())
    assert [row["sales_contract_no"] for row in rows] == ["053128", "053129", "053130"]
    assert [row["source_row_no"] for row in rows] == [4, 6, 8]
    assert {row["source_sheet_name"] for row in rows} == {"正单评审表"}
    first = rows[0]
    assert first["customer_po_no"] == "0006661356"
    assert first["customer_item_no"] == "0045841"
    assert first["customer_name"] == "WMC"
    assert first["export_country_code"] == ""  # Market and standards are not country codes.
    assert first["quantity"] == Decimal("3752")
    assert first["packing"] == "0/4"
    assert first["planned_inspection_date"] == "2026-09-15"
    assert first["shipment_date"] == "2026-09-21"
    assert first["source_cells"]["customer_po_no"] == {"cell": "E4", "value": "0006661356"}
    assert rows[1]["planned_inspection_date"] == "2026-09-17"
    assert rows[2]["customer_po_no"] == ""
    assert rows[2]["planned_inspection_date"] == "9/21-10% 10/23-80%"


@pytest.mark.parametrize("options,message", [
    ({"boundary": False}, "取消单分界"),
    ({"review_sheet": False}, "未找到正单评审表"),
    ({"order_type": False}, "缺少订单类型"),
])
def test_review_import_fails_closed_without_source_boundary_or_type(options, message):
    with pytest.raises(HTTPException) as error:
        parse(review_workbook(**options))
    assert error.value.status_code == 422
    assert message in error.value.detail


def test_review_preview_confirmation_preserves_scope_and_invalid_rows(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "qc_inspector")
        content = review_workbook()
        preview = client.post(
            "/api/qc-inspections/schedule-imports/preview",
            data={"factory_id": "huaxing", "week_key": "2026-W38", "request_id": "review-preview"},
            files={"file": ("BUZZ BEE.xlsx", content)},
        )
        assert preview.status_code == 201, preview.text
        batch = preview.json()
        assert batch["parser_version"] == "qc-review-unshipped-v1"
        assert batch["row_count"] == 3
        assert batch["blocking_count"] == 1
        assert [row["source_row_no"] for row in batch["rows"]] == [4, 6, 8]
        assert all(row["source_sheet_name"] == "正单评审表" for row in batch["rows"])
        invalid = batch["rows"][2]
        assert invalid["match_status"] == "INVALID"
        assert "PO不能为空" in invalid["validation_errors"]
        assert "计划验货日期格式必须为 YYYY-MM-DD" in invalid["validation_errors"]
        # Preview does not create orders; factory authorization remains enforced.
        orders = client.get("/api/qc-inspections/orders", params={"factory_id": "huaxing", "week": "2026-W38"})
        assert orders.json()["items"] == []
        denied = client.post(
            "/api/qc-inspections/schedule-imports/preview",
            data={"factory_id": "huakang-a", "week_key": "2026-W38", "request_id": "review-cross-factory"},
            files={"file": ("BUZZ BEE.xlsx", content)},
        )
        assert denied.status_code == 403
        confirmed = client.post(
            f"/api/qc-inspections/schedule-imports/{batch['id']}/confirm",
            json={"factory_id": "huaxing", "expected_revision": batch["revision"],
                  "request_id": "review-confirm", "decisions": [
                      {"row_id": row["id"], "action": "SKIP" if row["match_status"] == "INVALID" else "CREATE"}
                      for row in batch["rows"]
                  ]},
        )
        assert confirmed.status_code == 200, confirmed.text
        orders = client.get("/api/qc-inspections/orders", params={"factory_id": "huaxing", "week": "2026-W38"}).json()["items"]
        assert {row["sales_contract_no"] for row in orders} == {"053128", "053129"}


def test_selected_import_rows_remain_pending_and_resume_safely(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "qc_inspector")
        batch = client.post(
            "/api/qc-inspections/schedule-imports/preview",
            data={"factory_id": "huaxing", "week_key": "2026-W38", "request_id": "bulk-preview"},
            files={"file": ("BUZZ BEE.xlsx", review_workbook())},
        ).json()
        endpoint = f"/api/qc-inspections/schedule-imports/{batch['id']}"
        payload = {"factory_id": "huaxing", "expected_revision": batch["revision"],
                   "request_id": "bulk-first", "decisions": [{"row_id": batch["rows"][0]["id"], "action": "CREATE"}]}
        first = client.post(endpoint + "/confirm", json=payload)
        assert first.status_code == 200, first.text
        assert first.json()["status"] == "PARTIALLY_CONFIRMED"
        assert [row["decision_status"] for row in first.json()["rows"]] == ["CREATE", "PENDING", "PENDING"]
        assert client.post(endpoint + "/confirm", json=payload).json() == first.json()
        stale = {**payload, "request_id": "bulk-stale", "decisions": [{"row_id": batch["rows"][1]["id"], "action": "CREATE"}]}
        assert client.post(endpoint + "/confirm", json=stale).status_code == 409
        resumed = client.get(endpoint, params={"factory_id": "huaxing"}).json()
        invalid = {**stale, "request_id": "bulk-invalid", "expected_revision": resumed["revision"],
                   "decisions": [{"row_id": batch["rows"][2]["id"], "action": "CREATE"}]}
        assert client.post(endpoint + "/confirm", json=invalid).status_code == 422
        second = client.post(endpoint + "/confirm", json={**stale, "request_id": "bulk-second", "expected_revision": resumed["revision"]})
        assert second.status_code == 200, second.text
        assert [row["decision_status"] for row in second.json()["rows"]] == ["CREATE", "CREATE", "PENDING"]
        orders = client.get("/api/qc-inspections/orders", params={"factory_id": "huaxing"}).json()["items"]
        assert len(orders) == 2
        assert client.get(endpoint, params={"factory_id": "huakang-a"}).status_code == 403
