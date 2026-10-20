from copy import copy
from io import BytesIO

import openpyxl
import pytest
from openpyxl.worksheet.formula import ArrayFormula

from app.services import customer_order_unified as service
from app.services.huakang_a_order_legacy import unified_schedule as mapping


def template(customer):
    w = openpyxl.Workbook()
    w.remove(w.active)
    for name, marker in ((service.ITEM_SHEET, 8), (service.ORDER_SHEET, 6), (service.REVIEW_SHEET, 7)):
        s = w.create_sheet(name)
        headers = service.ITEM_HEADERS[:29] + mapping.ITEM_TAILS[customer] if name == service.ITEM_SHEET else service.ORDER_HEADERS
        for c, value in enumerate(headers, 1):
            s.cell(3, c, value)
        s.cell(marker, 1, '取消单')
        s.cell(marker + 2, 1, '已走货订单')
        s.cell(marker + 3, 1, '保留历史')
        s.cell(marker + 3, 2, f'=A{marker + 3}')
        s.merge_cells(start_row=marker, start_column=1, end_row=marker, end_column=3)
        if name == service.ITEM_SHEET:
            for c, value in {6: '7816-14', 7: 'GT', 8: 'BBJDTR-1850', 9: '翻斗车', 10: 'Dump Truck', 11: 2500, 12: 6, 20: '彩盒'}.items():
                s.cell(4, c, value)
            s.cell(4, 13, '=CEILING(K4/L4,1)')
        else:
            s.cell(4, 7, 'OLD')
            s.cell(4, 8, ArrayFormula(ref='H4', text='=IF(D4="","",ITEM表!I4)'))
            s.cell(4, 14, 11)
            s.cell(4, 15, '=J4*N4')
    buffer = BytesIO()
    w.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize('customer', mapping.ITEM_TAILS)
def test_common_headers_accept_different_customer_extensions(customer):
    content = template(customer)
    service.ensure_unified_schedule('new.xlsx', content, factory_id='huakang-a', customer_code=customer)
    service.ensure_unified_schedule('new.xlsx', content, factory_id='huaxing', customer_code='360')
    other = next(c for c in mapping.ITEM_TAILS if c != customer)
    service.ensure_unified_schedule('new.xlsx', content, factory_id='huakang-a', customer_code=other)


def test_custom_columns_optional_but_common_keys_required():
    w = openpyxl.load_workbook(BytesIO(template('360')))
    item = w[service.ITEM_SHEET]
    item.delete_cols(29, 30)
    for name in (service.ORDER_SHEET, service.REVIEW_SHEET):
        w[name]['T3'] = '跟客自定义备注'
    b = BytesIO()
    w.save(b)
    service.ensure_unified_schedule('new.xlsx', b.getvalue())
    item['H3'] = '错误公共字段'
    b = BytesIO()
    w.save(b)
    with pytest.raises(service.CustomerOrderUnifiedError, match='公共字段.*产品编号'):
        service.ensure_unified_schedule('new.xlsx', b.getvalue())


def test_reordered_extensions_write_by_header_and_preserve_unknown_cells():
    w = openpyxl.load_workbook(BytesIO(template('360')))
    s = w[service.ITEM_SHEET]
    for c in range(29, 39):
        s.cell(3, c).value = None
    for c, value in {29:'跟客负责人', 30:'备注', 31:'Port of Discharge', 34:'条码', 36:'Consolidated ID', 38:'MS Container Type (FCL/LCL)'}.items():
        s.cell(3, c).value = value
    s['AC5'], s['AD5'] = '人工跟客', '人工备注'
    mapping.write_extras(s, 5, {'barcode':'00123','port':'China','container_type':'LCL','consolidated_id':'FCL'}, '360')
    assert [s.cell(5,c).value for c in (31,34,36,38)] == ['China','00123','FCL','LCL']
    assert (s['AC5'].value,s['AD5'].value) == ('人工跟客','人工备注')
    # Missing or duplicate labels are not a reason to reject the common template.
    s['AI3'] = '条码'
    s['AH5'], s['AI5'] = '保留1', '保留2'
    mapping.write_extras(s, 5, {'barcode':'999'}, '360')
    assert (s['AH5'].value,s['AI5'].value) == ('保留1','保留2')


@pytest.mark.parametrize('customer', ['green-toys','headstart'])
def test_price_formulas_follow_reordered_headers(customer):
    w = openpyxl.load_workbook(BytesIO(template(customer)))
    s = w[service.ITEM_SHEET]
    s.insert_cols(30, 2)
    columns = mapping.extra_columns(s, customer)
    row = {'quantity':'25','unit_price_usd':'1.89','unit_price_hkd':'14.742','transportation_mode':"40'YT"}
    mapping.write_extras(s, 5, row, customer)
    assert s['AJ5'].value == 1.89
    assert 'AJ5*7.8' in s['AL5'].value
    assert 'K5*AJ5' in s['AK5'].value
    assert 'K5*AL5' in s['AM5'].value
    if customer == 'green-toys':
        assert columns['transportation_mode'] == 35
        assert s['AI5'].value == "40'YT"
        assert s['AR5'].value is None  # actual shipping mode


def test_actual_shipping_mode_not_written_when_planned_mode_removed():
    w = openpyxl.load_workbook(BytesIO(template('green-toys')))
    s = w[service.ITEM_SHEET]
    s['AG3'] = '人工备注'
    s['AP5'] = '实际运输方式'
    assert 'transportation_mode' not in mapping.extra_columns(s, 'green-toys')
    mapping.write_extras(s, 5, {'transportation_mode': "40'YT"}, 'green-toys')
    assert s['AP5'].value == '实际运输方式'


def test_green_history_identity_uses_base_po_and_item_even_below_cancel():
    content = template('green-toys')
    w = openpyxl.load_workbook(BytesIO(content))
    s = w[service.ITEM_SHEET]
    s['F15'], s['H15'], s['K15'] = '7816-2', 'BOOK-1', 50
    b = BytesIO()
    w.save(b)
    history = service.read_unified_history(b.getvalue(), factory_id='huakang-a', customer_code='green-toys')
    assert {h.reference_no for h in history} == {'7816-14', '7816-2'}
    record = {'contract_no': '7816', 'customer_po': '7816-1', 'item_no': 'BBJDTR-1850'}
    assigned = {}
    mapping.reconcile_green_po(record, history, assigned)
    assert record['customer_po'] == '7816-14'
    new = {'contract_no': '7816', 'item_no': 'NEW'}
    mapping.reconcile_green_po(new, history, assigned)
    assert new['customer_po'] == '7816-15'
    assert mapping.identity('headstart', '2022116AAA', '25833−0206') == mapping.identity('headstart', '2022116AA', '25833-0206')


@pytest.mark.parametrize('customer', mapping.ITEM_TAILS)
def test_independent_append_preserves_history_and_fills_reserved_array_formulas(customer):
    w = openpyxl.load_workbook(BytesIO(template(customer)))
    before = {s.title: service._marker_row(s) for s in w}
    slots = mapping.output_slots(w, 6)
    assert all(len(rows) == 6 for rows in slots.values())
    for s in w:
        marker = service._marker_row(s)
        delta = marker - before[s.title]
        assert s.cell(before[s.title] + 3 + delta, 1).value == '保留历史'
        assert s.cell(before[s.title] + 3 + delta, 2).value == f'=A{before[s.title] + 3 + delta}'
        assert f'A{marker}:C{marker}' in {str(r) for r in s.merged_cells.ranges}
        assert max(slots[s.title]) < marker
        if s.title != service.ITEM_SHEET:
            assert s['N4'].value == 11
            for r in slots[s.title]:
                assert isinstance(s.cell(r, 8).value, ArrayFormula)
                assert s.cell(r, 8).value.ref == f'H{r}'
                assert f'D{r}' in s.cell(r, 8).value.text
                assert s.cell(r, 14).value is None
    b = BytesIO()
    w.save(b)
    loaded = openpyxl.load_workbook(BytesIO(b.getvalue()))
    assert isinstance(loaded[service.ORDER_SHEET]['H5'].value, ArrayFormula)


@pytest.mark.parametrize('customer', mapping.ITEM_TAILS)
def test_preview_and_export_customer_fields(monkeypatch, customer):
    # API regression fixtures reload service modules between tests.
    from app.services.huakang_a_order_legacy import schedule_parser, green_toys_headstart
    record = {
        'contract_no': 'RL-765431-19' if customer == '360' else '2022802AB',
        'customer_release': 'RL765432118', 'customer_po': '7654321',
        'item_no': 'NEW', 'description': 'New toy', 'quantity': 2500,
        'master_carton_qty': 6, 'barcode': '001234567890', 'container_type': 'LCL',
        'consolidated_id': 'FCL', 'port': 'China', 'unit_price_usd': 1.89,
        'country': '美国', 'product_packaging': 'Window Box', 'customer_name': 'Example Retailer',
        'planned_inspection_date': '2026-09-20', 'factory_commit_date': '2026-09-21',
        'parse_ok': True, 'parse_warnings': [],
    }
    if customer == '360':
        record['outer_carton_sticker'] = 'Apply shipping label to each master carton'
        record['outer_carton_sticker_sources'] = [{
            'file_name': 'po.pdf', 'relative_path': 'po.pdf', 'page': 2,
            'raw_text': 'Apply shipping label to each master carton',
        }]
        monkeypatch.setattr(schedule_parser, 'parse_po_file', lambda *_: copy(record))
    else:
        fn = 'parse_green_toys_image' if customer == 'green-toys' else 'parse_headstart_pdf'
        monkeypatch.setattr(green_toys_headstart, fn, lambda *_: [copy(record)])
    output, _, preview = service.export_unified_customer_schedule(
        customer_code=customer, factory_id='huakang-a', received_date='2026-09-07',
        po_files=[('po.pdf' if customer != 'green-toys' else 'po.png', b'po')],
        schedule_file_name='new.xlsx', schedule_content=template(customer),
    )
    row = preview['rows'][0]
    assert row['standard'] == ''
    assert row['packaging'] == 'Window Box'
    assert row['carton_count'] == '417'
    w = openpyxl.load_workbook(BytesIO(output))
    s = w[service.ITEM_SHEET]
    item_row = next(c.row for c in s._cells.values() if c.column == 8 and c.value == 'NEW')
    assert s.cell(item_row, 13).value == f'=CEILING(K{item_row}/L{item_row},1)'
    assert s.cell(item_row, 14).value is None
    if customer == '360':
        assert s.cell(item_row, 3).value == 'RL765432118'
        assert s.cell(item_row, 4).value == s.cell(item_row, 5).value == 'RL-765431-19'
        assert s.cell(item_row, 6).value == '7654321'
        assert row['customer_release_no'] == 'RL765432118'
        assert row['lineage']['production_no'] == 'po.pdf · Customer Release No.（客户 Release/OH 号）'
        assert row['lineage']['po_no'] == 'po.pdf · Customer PO Number'
        assert row['customer_label'] == 'Apply shipping label to each master carton'
        assert row['lineage']['customer_label'] == 'po.pdf · 第2页 · Apply shipping label to each master carton'
        assert s.cell(item_row, 17).value == 'Apply shipping label to each master carton'
        assert [s.cell(item_row, c).value for c in range(30, 34)] == ['001234567890', 'LCL', 'FCL', 'China']
    else:
        assert s['AH5'].value == 1.89
        assert '7.8' in s['AJ5'].value
        assert 'ROUND' not in s['AI5'].value
        assert s['AL5'].value is None
    for name in (service.ORDER_SHEET, service.REVIEW_SHEET):
        summary_row = next(c.row for c in w[name]._cells.values() if c.column == 7 and c.value == 'NEW')
        assert [w[name].cell(summary_row, c).value for c in range(3, 8)] == [s.cell(item_row, c).value for c in range(4, 9)]


def _release_record(reference, customer_release, customer_po='7654321'):
    return {
        'contract_no': reference, 'customer_release': customer_release,
        'customer_po': customer_po, 'item_no': 'NEW', 'description': 'New toy',
        'quantity': 120, 'master_carton_qty': 6, 'customer_name': 'Example Retailer',
        'planned_inspection_date': '2026-09-20', 'factory_commit_date': '2026-09-21',
        'parse_ok': True, 'parse_warnings': [],
    }


def test_360_same_customer_po_keeps_distinct_rl_orders_and_preserves_history(monkeypatch):
    from app.services.huakang_a_order_legacy import schedule_parser
    records = {
        'first.pdf': _release_record('RL-765431-19', 'OH-001'),
        'second.pdf': _release_record('RL-765431-20', 'OH-002'),
    }
    monkeypatch.setattr(schedule_parser, 'parse_po_file',
                        lambda path, *_: copy(records[path.name.split('_', 1)[1]]))
    original = openpyxl.load_workbook(BytesIO(template('360')))
    old = original[service.ITEM_SHEET]
    old_product = old.cell(4, 8).value
    # Previous generated/manual identifiers remain source evidence, not a
    # migration target when exporting newly received releases.
    old_values = ['RL-765431-18', 'OH-OLD', 'RL-765431-18', 'RL-765431-18']
    for column, value in enumerate(old_values, 3):
        old.cell(4, column, value)
    original_content = BytesIO()
    original.save(original_content)
    output, _, preview = service.export_unified_customer_schedule(
        customer_code='360', factory_id='huakang-a', received_date='2026-09-07',
        po_files=[(name, name.encode()) for name in records],
        schedule_file_name='new.xlsx', schedule_content=original_content.getvalue(),
    )
    assert preview['summary']['blocked'] == 0
    assert {row['reference_no'] for row in preview['rows']} == {'RL-765431-19', 'RL-765431-20'}
    workbook = openpyxl.load_workbook(BytesIO(output))
    item = workbook[service.ITEM_SHEET]
    old_row = next(c.row for c in item._cells.values() if c.column == 8 and c.value == old_product)
    assert [item.cell(old_row, c).value for c in range(3, 7)] == old_values
    for record in records.values():
        item_row = next(c.row for c in item._cells.values()
                        if c.column == 5 and c.value == record['contract_no'])
        assert [item.cell(item_row, c).value for c in range(3, 7)] == [
            record['customer_release'], record['contract_no'], record['contract_no'], '7654321',
        ]
        for name in (service.ORDER_SHEET, service.REVIEW_SHEET):
            summary = workbook[name]
            summary_row = next(c.row for c in summary._cells.values()
                               if c.column == 4 and c.value == record['contract_no'])
            assert [summary.cell(summary_row, c).value for c in range(3, 6)] == [
                record['contract_no'], record['contract_no'], '7654321',
            ]


def test_360_different_oh_does_not_create_a_new_identity_for_same_rl(monkeypatch):
    from app.services.huakang_a_order_legacy import schedule_parser
    records = {
        'first.pdf': _release_record('RL-765431-19', 'OH-001'),
        'second.pdf': _release_record('RL-765431-19', 'OH-002'),
    }
    monkeypatch.setattr(schedule_parser, 'parse_po_file',
                        lambda path, *_: copy(records[path.name.split('_', 1)[1]]))
    preview = service.create_unified_customer_preview(
        customer_code='360', factory_id='huakang-a', received_date='2026-09-07',
        po_files=[(name, name.encode()) for name in records],
        schedule_file_name='new.xlsx', schedule_content=template('360'),
    )
    assert {row['reference_no'] for row in preview['rows']} == {'RL-765431-19'}
    assert {row['production_no'] for row in preview['rows']} == {'OH-001', 'OH-002'}
    assert any(issue['code'] == 'duplicate_batch_order_line'
               for row in preview['rows'] for issue in row['issues'])


def test_360_missing_customer_identifiers_are_not_replaced_by_rl():
    row = mapping.map_record({}, _release_record('RL-765431-19', '', ''), '360')
    assert row['contract_no'] == row['reference_no'] == 'RL-765431-19'
    assert row['production_no'] == row['po_no'] == row['customer_release_no'] == ''
    assert 'customer_label' not in row['lineage']


def test_reserved_summary_formulas_take_priority_over_legacy_contract_formulas():
    w = openpyxl.load_workbook(BytesIO(template('360')))
    s = w[service.REVIEW_SHEET]
    # The latest 360 file has a complete SO-based formula row after a blank gap.
    for c in (1, 2, 8, 9, 10, 11, 12, 13, 15, 16, 17, 18):
        s.cell(6, c, '=IF(D6="","",ITEM表!E6)')
    slots = mapping.output_slots(w, 4)
    assert slots[service.REVIEW_SHEET][0] == 5
    assert s['H5'].value.startswith('=IF(D5=')
    assert s['H6'].value == '=IF(D6="","",ITEM表!E6)'
    assert s['N5'].value is None
