import hashlib
import importlib
import json
from io import BytesIO

from openpyxl import load_workbook

from test_internal_quote_api import ALL_SECTION_CODES, create_payload, login, logout, make_client


def mark_all_sections_not_applicable(quote_id: str) -> None:
    db_module = importlib.import_module("app.db")
    quote_models = importlib.import_module("app.models.internal_quote")
    with db_module.SessionLocal() as db:
        quote = db.get(quote_models.InternalQuote, quote_id)
        sections = db.query(quote_models.InternalQuoteSection).filter_by(quote_id=quote_id).all()
        for section in sections:
            section.status = "not_applicable"
            section.calculation_status = "not_applicable"
            section.dependency_status = "current"
            if section.department == "sales":
                section.submitted_by_id = quote.created_by
                section.submitted_by = quote.created_by_name
        quote.status = "ready_for_final_review"
        db.commit()


def mark_required_sections_not_applicable(quote_id: str) -> None:
    db_module = importlib.import_module("app.db")
    quote_models = importlib.import_module("app.models.internal_quote")
    with db_module.SessionLocal() as db:
        quote = db.get(quote_models.InternalQuote, quote_id)
        sections = db.query(quote_models.InternalQuoteSection).filter_by(quote_id=quote_id).all()
        for section in sections:
            if not section.is_required:
                continue
            section.status = "not_applicable"
            section.calculation_status = "not_applicable"
            section.dependency_status = "current"
            if section.department == "sales":
                section.submitted_by_id = quote.created_by
                section.submitted_by = quote.created_by_name
        quote.status = "ready_for_final_review"
        db.commit()


def mark_legacy_final_pending(quote_id: str) -> None:
    db_module = importlib.import_module("app.db")
    quote_models = importlib.import_module("app.models.internal_quote")
    calculator = importlib.import_module("app.services.internal_quote_calculator")
    release_service = importlib.import_module("app.services.internal_quote_release")
    with db_module.SessionLocal() as db:
        quote = db.get(quote_models.InternalQuote, quote_id)
        sections = release_service._quote_sections(db, quote_id)
        quote.header_revision += 1
        manifest = release_service._release_manifest(quote, sections)
        quote.status = "final_reviewing"
        quote.final_release_status = "pending"
        quote.final_submission_revision += 1
        quote.final_submission_manifest_json = calculator.canonical_json(manifest)
        quote.final_submitted_by = quote.created_by
        quote.final_submitted_by_name = quote.created_by_name
        quote.final_submitted_at = quote.updated_at
        db.commit()


def test_p4_release_ignores_inactive_optional_sections_and_marks_them_in_structured_data(monkeypatch):
    with make_client(monkeypatch) as client:
        submitter = login(
            client,
            "iq_p4_optional_submitter",
            "sales_customer_owner",
            "sales-business",
        )
        payload = create_payload(suffix="P4-OPTIONAL")
        payload["business_owner_id"] = submitter["id"]
        payload["business_owner_name"] = submitter["display_name"]
        created = client.post("/api/internal-quotes", json=payload)
        assert created.status_code == 201, created.text
        quote_id = created.json()["id"]
        mark_required_sections_not_applicable(quote_id)

        submitted = client.post(
            f"/api/internal-quotes/{quote_id}/final-submit",
            json={"revision": 1},
        )
        assert submitted.status_code == 200, submitted.text
        export = submitted.json()["export"]
        downloaded = client.get(
            f"/api/internal-quotes/{quote_id}/exports/{export['id']}/download"
        )
        workbook = load_workbook(BytesIO(downloaded.content), read_only=True, data_only=False)
        rows = list(workbook["结构化数据"].iter_rows(min_row=4, values_only=True))
        payload_rows = {row[1]: row for row in rows if row[0] == "payload"}
        assert payload_rows["sales"][3] == "not_applicable"
        assert payload_rows["sales"][11] == "是"
        assert payload_rows["electronic"][3] == "draft"
        assert payload_rows["electronic"][11] == "否"
        workbook.close()


def test_p4_responsible_sales_followup_releases_once_and_hands_off_final_artifact(monkeypatch):
    with make_client(monkeypatch) as client:
        submitter = login(
            client,
            "iq_p4_submitter",
            "sales_customer_owner",
            "sales-business",
        )
        assert "internal_quote:final_submit" in submitter["permissions"]
        assert "internal_quote:final_approve" not in submitter["permissions"]
        payload = create_payload(suffix="P4-RELEASE", participating_sections=ALL_SECTION_CODES)
        payload["business_owner_id"] = submitter["id"]
        payload["business_owner_name"] = submitter["display_name"]
        quote = client.post("/api/internal-quotes", json=payload).json()
        quote_id = quote["id"]

        blocked = client.post(
            f"/api/internal-quotes/{quote_id}/final-submit",
            json={"revision": 1},
        )
        assert blocked.status_code == 409

        mark_all_sections_not_applicable(quote_id)
        p3_export = client.post(f"/api/internal-quotes/{quote_id}/exports")
        assert p3_export.status_code == 201
        assert p3_export.json()["release_stage"] == "p3_section_approved"
        assert client.get(
            "/api/customer-price/internal-quote-artifacts?factory_id=huaxing"
        ).json() == []

        logout(client)
        login(client, "iq_p4_unrelated_owner", "sales_customer_owner", "sales-business")
        unrelated_release = client.post(
            f"/api/internal-quotes/{quote_id}/final-submit",
            json={"revision": 1},
        )
        assert unrelated_release.status_code == 403
        assert "负责本单的业务跟客" in unrelated_release.json()["detail"]

        logout(client)
        login(
            client,
            "iq_p4_submitter",
            "sales_customer_owner",
            "sales-business",
        )
        released = client.post(
            f"/api/internal-quotes/{quote_id}/final-submit",
            json={"revision": 1},
        )
        assert released.status_code == 200, released.text
        result = released.json()
        assert result["quote"]["status"] == "exported"
        assert result["quote"]["header_revision"] == 3
        assert result["quote"]["final_release_status"] == "approved"
        assert result["quote"]["final_release_revision"] == 1
        assert result["review"]["decision"] == "approve"
        assert result["review"]["submitted_by"] == submitter["id"]
        assert result["review"]["actor_id"] == submitter["id"]
        notifications_after_review = client.get("/api/system/notifications").json()
        assert not any(
            item["payload"].get("event") == "final_release_submitted"
            and item["payload"].get("quote_id") == quote_id
            for item in notifications_after_review
        )
        final_notification = next(
            item
            for item in notifications_after_review
            if item["payload"].get("event") == "final_release_approved"
            and item["payload"].get("quote_id") == quote_id
        )
        assert final_notification["payload"]["route"] == (
            f"/modules/sales-business/internal-quote-desk/{quote_id}/summary?factory=huaxing"
        )
        artifact_notification = next(
            item
            for item in notifications_after_review
            if item["payload"].get("event") == "customer_price_artifact_available"
            and item["payload"].get("quote_id") == quote_id
        )
        assert artifact_notification["status"] == "unread"
        final_export = result["export"]
        assert final_export["template_version"] == "internal-quote-p4-v2"
        assert final_export["release_stage"] == "p4_final_approved"
        assert final_export["export_manifest"]["workbook_layout_version"] == "internal-quote-unified-desk-v3"
        assert final_export["export_manifest"]["p4_final_release_required"] is False
        assert final_export["export_manifest"]["final_reviewed_by"] == submitter["id"]

        export_download = client.get(
            f"/api/internal-quotes/{quote_id}/exports/{final_export['id']}/download"
        )
        assert export_download.status_code == 200
        workbook = load_workbook(BytesIO(export_download.content), read_only=True, data_only=False)
        quote_sheet = workbook["报价明细"]
        assert quote_sheet.sheet_state == "visible"
        assert quote_sheet["A8"].value == "内部报价测试产品报价"
        assert quote_sheet["A8"].font.name == "宋体"
        assert quote_sheet["A8"].font.sz == 14
        assert quote_sheet["D2"].number_format == "0.0"
        assert quote_sheet["N2"].number_format == "0.00%"
        assert quote_sheet["D6"].number_format == "#,##0"
        assert quote_sheet["N6"].number_format == "0.00"
        assert quote_sheet["C9"].fill.fill_type is None
        assert quote_sheet["B10"].value is None
        assert quote_sheet["C41"].value == "报客价："
        assert quote_sheet["D42"].value == 3.5
        assert quote_sheet["C46"].value == "旺季价"
        assert all(workbook[name].sheet_state == "veryHidden" for name in workbook.sheetnames[1:])
        assert workbook["审批与版本"]["B4"].value == "P4 最终业务放行"
        assert workbook["审批与版本"]["B5"].value == "最终业务放行完成，可交接客价转换台"
        structured = workbook["结构化数据"]
        assert structured["B2"].value == "internal-quote-structured-data-v1"
        structured_rows = list(structured.iter_rows(min_row=4, values_only=True))
        assert {
            (record_type, department)
            for record_type, department in {(row[0], row[1]) for row in structured_rows}
            if department != "quote"
        } == {
            (record_type, department)
            for record_type in ("payload", "calculation")
            for department in (
                "sales",
                "engineering",
                "electronic",
                "molding",
                "painting",
                "slush",
                "sewing",
                "assembly",
            )
        }
        snapshot_parts = sorted(
            (row for row in structured_rows if row[0] == "reference_snapshot" and row[1] == "quote"),
            key=lambda row: row[8],
        )
        snapshot = json.loads("".join(str(row[10]) for row in snapshot_parts))
        assert snapshot["machine_prices"]
        assert snapshot["fx"]["hkd_usd"] == "7.8"
        for record_type in ("payload", "calculation"):
            sales_parts = sorted(
                (row for row in structured_rows if row[0] == record_type and row[1] == "sales"),
                key=lambda row: row[8],
            )
            assert len(sales_parts) == sales_parts[0][9]
            assert json.loads("".join(str(row[10]) for row in sales_parts)) == {}
        workbook.close()

        artifacts = client.get(
            "/api/customer-price/internal-quote-artifacts?factory_id=huaxing"
        )
        assert artifacts.status_code == 200, artifacts.text
        assert len(artifacts.json()) == 1
        handoff = artifacts.json()[0]
        assert handoff["export_id"] == final_export["id"]
        assert handoff["status"] == "available"
        assert handoff["artifact_manifest"]["release_revision"] == 1
        assert handoff["artifact_manifest"]["template_version"] == "internal-quote-p4-v2"
        assert handoff["artifact_manifest"]["workbook_layout_version"] == "internal-quote-unified-desk-v3"
        assert handoff["sha256"] == final_export["sha256"]

        artifact_download = client.get(
            f"/api/customer-price/internal-quote-artifacts/{handoff['id']}/download"
        )
        assert artifact_download.status_code == 200
        assert artifact_download.headers["x-internal-quote-release-stage"] == "p4_final_approved"
        assert hashlib.sha256(artifact_download.content).hexdigest() == handoff["sha256"]

        consumed = client.post(
            f"/api/customer-price/internal-quote-artifacts/{handoff['id']}/consume",
            json={"consumer_reference": "customer-conversion-001"},
        )
        assert consumed.status_code == 200
        assert consumed.json()["status"] == "consumed"
        assert consumed.json()["consumer_reference"] == "customer-conversion-001"
        notifications_after_consume = client.get("/api/system/notifications").json()
        assert next(
            item for item in notifications_after_consume if item["id"] == artifact_notification["id"]
        )["status"] == "handled"
        idempotent = client.post(
            f"/api/customer-price/internal-quote-artifacts/{handoff['id']}/consume",
            json={"consumer_reference": "customer-conversion-001"},
        )
        assert idempotent.status_code == 200
        assert idempotent.json()["id"] == handoff["id"]
        assert idempotent.json()["consumer_reference"] == "customer-conversion-001"
        duplicate = client.post(
            f"/api/customer-price/internal-quote-artifacts/{handoff['id']}/consume",
            json={"consumer_reference": "customer-conversion-002"},
        )
        assert duplicate.status_code == 409

        repeated_export = client.post(f"/api/internal-quotes/{quote_id}/exports")
        assert repeated_export.status_code == 201
        assert repeated_export.json()["id"] == final_export["id"]
        assert client.get(
            "/api/customer-price/internal-quote-artifacts?factory_id=huaxing"
        ).json() == []
        consumed_artifacts = client.get(
            "/api/customer-price/internal-quote-artifacts?factory_id=huaxing&status=consumed"
        )
        assert [item["id"] for item in consumed_artifacts.json()] == [handoff["id"]]

        db_module = importlib.import_module("app.db")
        auth_models = importlib.import_module("app.models.auth")
        with db_module.SessionLocal() as db:
            stale_artifact_notification = db.get(
                auth_models.SystemNotification,
                artifact_notification["id"],
            )
            stale_artifact_notification.status = "unread"
            stale_artifact_notification.read_at = ""
            stale_artifact_notification.handled_at = ""
            db.commit()

        reopened = client.post(
            f"/api/internal-quotes/{quote_id}/sections/sales/reopen",
            json={"revision": 1, "reason": "客户数量变更，重新核价"},
        )
        assert reopened.status_code == 200, reopened.text
        detail = client.get(f"/api/internal-quotes/{quote_id}").json()
        assert detail["status"] == "drafting"
        assert detail["final_release_status"] == "invalidated"
        notifications_after_reopen = client.get("/api/system/notifications").json()
        assert next(
            item for item in notifications_after_reopen if item["id"] == artifact_notification["id"]
        )["status"] == "handled"
        revoked = client.get(
            "/api/customer-price/internal-quote-artifacts?factory_id=huaxing&status=revoked"
        )
        assert revoked.status_code == 200
        assert revoked.json()[0]["id"] == handoff["id"]
        assert revoked.json()[0]["revoke_reason"]


def test_p4_legacy_layout_export_is_refreshed_without_replacing_release_handoff(monkeypatch):
    with make_client(monkeypatch) as client:
        submitter = login(
            client,
            "iq_p4_layout_refresh_submitter",
            "sales_customer_owner",
            "sales-business",
        )
        payload = create_payload(suffix="P4-LAYOUT-REFRESH", participating_sections=ALL_SECTION_CODES)
        payload["business_owner_id"] = submitter["id"]
        payload["business_owner_name"] = submitter["display_name"]
        quote = client.post("/api/internal-quotes", json=payload).json()
        quote_id = quote["id"]
        mark_all_sections_not_applicable(quote_id)
        submitted = client.post(
            f"/api/internal-quotes/{quote_id}/final-submit",
            json={"revision": 1},
        )
        assert submitted.status_code == 200, submitted.text
        legacy_export = submitted.json()["export"]

        db_module = importlib.import_module("app.db")
        quote_models = importlib.import_module("app.models.internal_quote")
        with db_module.SessionLocal() as db:
            export_record = db.get(quote_models.InternalQuoteExportFile, legacy_export["id"])
            legacy_manifest = json.loads(export_record.export_manifest_json)
            legacy_manifest.pop("workbook_layout_version", None)
            export_record.export_manifest_json = json.dumps(
                legacy_manifest,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            )
            handoff_record = db.query(quote_models.InternalQuoteArtifactHandoff).filter_by(
                quote_id=quote_id,
                release_revision=1,
            ).one()
            legacy_handoff_id = handoff_record.id
            legacy_handoff_manifest = json.loads(handoff_record.artifact_manifest_json)
            legacy_handoff_manifest.pop("workbook_layout_version", None)
            handoff_record.artifact_manifest_json = json.dumps(
                legacy_handoff_manifest,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            )
            db.commit()

        refreshed_response = client.post(f"/api/internal-quotes/{quote_id}/exports")
        assert refreshed_response.status_code == 201, refreshed_response.text
        refreshed = refreshed_response.json()
        assert refreshed["id"] != legacy_export["id"]
        assert refreshed["template_version"] == "internal-quote-p4-v2"
        assert refreshed["export_manifest"]["workbook_layout_version"] == "internal-quote-unified-desk-v3"

        downloaded = client.get(
            f"/api/internal-quotes/{quote_id}/exports/{refreshed['id']}/download"
        )
        assert downloaded.status_code == 200
        workbook = load_workbook(BytesIO(downloaded.content), read_only=True, data_only=False)
        assert workbook["报价明细"]["A8"].value == "内部报价测试产品报价"
        assert workbook["报价明细"]["C41"].value == "报客价："
        assert all(workbook[name].sheet_state == "veryHidden" for name in workbook.sheetnames[1:])
        workbook.close()

        repeated = client.post(f"/api/internal-quotes/{quote_id}/exports")
        assert repeated.status_code == 201
        assert repeated.json()["id"] == refreshed["id"]

        history = client.get(f"/api/internal-quotes/{quote_id}/exports").json()
        history_by_id = {item["id"]: item for item in history}
        assert history_by_id[legacy_export["id"]]["status"] == "superseded"
        assert history_by_id[refreshed["id"]]["status"] == "current"

        artifacts = client.get(
            "/api/customer-price/internal-quote-artifacts?factory_id=huaxing"
        )
        assert artifacts.status_code == 200, artifacts.text
        assert len(artifacts.json()) == 1
        preserved_handoff = artifacts.json()[0]
        assert preserved_handoff["id"] == legacy_handoff_id
        assert preserved_handoff["export_id"] == refreshed["id"]
        assert preserved_handoff["status"] == "available"
        assert preserved_handoff["artifact_manifest"]["workbook_layout_version"] == "internal-quote-unified-desk-v3"


def test_p4_final_rejection_is_immutable_and_can_be_resubmitted(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "iq_p4_reject_submitter", "sales_customer_supervisor", "sales-business")
        quote = client.post(
            "/api/internal-quotes",
            json=create_payload(suffix="P4-REJECT", participating_sections=ALL_SECTION_CODES),
        ).json()
        quote_id = quote["id"]
        mark_all_sections_not_applicable(quote_id)
        mark_legacy_final_pending(quote_id)

        logout(client)
        login(client, "iq_p4_reject_reviewer", "sales_customer_supervisor", "sales-business")
        missing_reason = client.post(
            f"/api/internal-quotes/{quote_id}/final-review",
            json={"revision": 2, "decision": "reject", "reason": ""},
        )
        assert missing_reason.status_code == 422
        rejected = client.post(
            f"/api/internal-quotes/{quote_id}/final-review",
            json={"revision": 2, "decision": "reject", "reason": "运费场景需业务确认"},
        )
        assert rejected.status_code == 200
        assert rejected.json()["quote"]["status"] == "ready_for_final_review"
        assert rejected.json()["quote"]["header_revision"] == 3
        assert rejected.json()["quote"]["final_release_status"] == "rejected"
        assert rejected.json()["export"] is None

        logout(client)
        login(client, "iq_p4_reject_submitter", "sales_customer_supervisor", "sales-business")
        resubmitted = client.post(
            f"/api/internal-quotes/{quote_id}/final-submit",
            json={"revision": 3},
        )
        assert resubmitted.status_code == 200
        assert resubmitted.json()["quote"]["final_submission_revision"] == 2
        assert resubmitted.json()["quote"]["final_release_status"] == "approved"

        reviews = client.get(f"/api/internal-quotes/{quote_id}/final-reviews")
        assert reviews.status_code == 200
        assert [(item["submission_revision"], item["decision"]) for item in reviews.json()] == [
            (2, "approve"),
            (1, "reject")
        ]


def test_p4_existing_pending_release_can_be_completed_by_responsible_followup(monkeypatch):
    with make_client(monkeypatch) as client:
        submitter = login(
            client,
            "iq_p4_pending_followup",
            "sales_customer_owner",
            "sales-business",
        )
        quote = client.post(
            "/api/internal-quotes",
            json=create_payload(suffix="P4-SELF-RETURN", participating_sections=ALL_SECTION_CODES),
        ).json()
        quote_id = quote["id"]
        mark_all_sections_not_applicable(quote_id)
        mark_legacy_final_pending(quote_id)

        logout(client)
        login(client, "iq_p4_pending_other", "sales_customer_owner", "sales-business")
        unrelated = client.post(
            f"/api/internal-quotes/{quote_id}/final-review",
            json={"revision": 2, "decision": "approve", "reason": "非负责跟客"},
        )
        assert unrelated.status_code == 403
        assert "负责本单的业务跟客" in unrelated.json()["detail"]

        logout(client)
        login(client, "iq_p4_pending_followup", "sales_customer_owner", "sales-business")
        released = client.post(
            f"/api/internal-quotes/{quote_id}/final-review",
            json={"revision": 2, "decision": "approve", "reason": "负责跟客确认放行"},
        )
        assert released.status_code == 200, released.text
        released_body = released.json()
        assert released_body["quote"]["status"] == "exported"
        assert released_body["quote"]["header_revision"] == 3
        assert released_body["quote"]["final_release_status"] == "approved"
        assert released_body["review"]["decision"] == "approve"
        assert released_body["review"]["submitted_by"] == submitter["id"]
        assert released_body["review"]["actor_id"] == submitter["id"]


def test_p4_quote_version_comparison_uses_clone_lineage_and_section_snapshots(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "iq_p4_compare", "sales_customer_supervisor", "sales-business")
        base = client.post(
            "/api/internal-quotes",
            json=create_payload(suffix="P4-COMPARE-BASE", participating_sections=ALL_SECTION_CODES),
        ).json()
        target = client.post(
            f"/api/internal-quotes/{base['id']}/clone",
            json={
                "quote_no": "IQ-TEST-P4-COMPARE-TARGET",
                "version_label": "V2",
                "business_owner_id": "owner-compare",
                "business_owner_name": "版本负责人",
                "target_date": "2026-09-01",
            },
        ).json()
        unrelated = client.post(
            "/api/internal-quotes",
            json=create_payload(suffix="P4-COMPARE-OTHER", participating_sections=ALL_SECTION_CODES),
        ).json()

        db_module = importlib.import_module("app.db")
        quote_models = importlib.import_module("app.models.internal_quote")
        with db_module.SessionLocal() as db:
            base_molding = db.query(quote_models.InternalQuoteSection).filter_by(
                quote_id=base["id"], department="molding"
            ).one()
            target_molding = db.query(quote_models.InternalQuoteSection).filter_by(
                quote_id=target["id"], department="molding"
            ).one()
            base_molding.payload_json = json.dumps({"machine_cost_hkd": "10.0000"})
            base_molding.calculation_json = json.dumps({"totals": {"total_hkd": "10.0000"}})
            base_molding.calculation_status = "valid"
            target_molding.payload_json = json.dumps({"machine_cost_hkd": "12.5000"})
            target_molding.calculation_json = json.dumps({"totals": {"total_hkd": "12.5000"}})
            target_molding.calculation_status = "valid"
            target_molding.revision = 2
            db.commit()

        candidates = client.get(f"/api/internal-quotes/{target['id']}/version-candidates")
        assert candidates.status_code == 200
        assert [item["id"] for item in candidates.json()] == [base["id"]]

        comparison = client.get(
            f"/api/internal-quotes/{target['id']}/compare/{base['id']}"
        )
        assert comparison.status_code == 200, comparison.text
        result = comparison.json()
        assert result["base"]["id"] == base["id"]
        assert result["target"]["id"] == target["id"]
        assert result["total_before_hkd"] == "10.0000"
        assert result["total_after_hkd"] == "12.5000"
        assert result["total_delta_hkd"] == "2.5000"
        molding = next(item for item in result["sections"] if item["section_code"] == "molding")
        assert molding["before_revision"] == 1
        assert molding["after_revision"] == 2
        assert molding["payload_changes"] == [
            {"path": "machine_cost_hkd", "before": "10.0000", "after": "12.5000"}
        ]

        invalid = client.get(
            f"/api/internal-quotes/{target['id']}/compare/{unrelated['id']}"
        )
        assert invalid.status_code == 400
