import io
import json
import hashlib
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from openpyxl import Workbook
from pypdf import PdfReader

from app.schemas.carton_supplier_portal import MarkLayoutConfig
from app.services.carton_mark_zuru import extract_zuru, prepare_records
from test_carton_mark_api import _pdf_bytes


def excel(profile="esal", multiple=False, blank_template=False):
    book = Workbook(); sheet = book.active; sheet.title = "SM"
    if profile == "zumiez":
        for row, value in enumerate(("ZUMIEZ INC.", "PO NUMBER:00484860", "ZUMIEZ ITEM NO.: 00415647",
                "ITEM DESCRIPTION: CURRENT PRODUCT", "QUANTITY:", "CTN QTY:", "COUNTRY OF ORIGIN: CHINA"), 1):
            sheet.cell(row, 1, value)
    else:
        for start, item in ((1, "00123"), (22, "00999")) if multiple else ((1, "00123"),):
            for offset, (name, value) in enumerate((("Item No:", item), ("Product Description:", "CURRENT PRODUCT " + item),
                    ("Barcode:", "00193052021172"), ("Brand:", "ZURU PETS ALIVE"), ("Qty/Ctn:", "6 PCS"),
                    ("Carton Net Weight:", "5.04 KGS"), ("Carton Gross Weight:", "5.93 KGS"),
                    ("Carton Size(cm):", "56.5 x 31.4 x 63.5 CM"), ("Carton CBM:", "0.113 CBM"), ("PI No.", "001100036670"))):
                sheet.cell(start+offset, 1, name); sheet.cell(start+offset, 2, value)
        if blank_template:
            sheet.cell(45,1,"Item No:");sheet.cell(45,2,"#N/A")
            sheet.cell(46,1,"Product Description:");sheet.cell(46,2,"#N/A")
    stream=io.BytesIO();book.save(stream);book.close();return stream.getvalue()


def original(profile="esal", **kwargs):
    region=dict(x=10,y=10,width=20,height=20)
    result=dict(excel_file_name="source.xlsx",excel_bytes=excel(profile),contract_no="4500215651",item_no="00123",
        layout=dict(name="ZURU",version=1,reference_bytes=_pdf_bytes(),config=MarkLayoutConfig(series="zuru",
            frame_style="original_red",logo_region=region,front_artwork_region=region,side_artwork_region=region,origin_artwork_region=region).model_dump()))
    result.update(kwargs);return result


def test_multi_item_blocks_preserve_leading_zero_skip_empty_lookups_and_never_mix_items():
    content=excel(multiple=True,blank_template=True)
    records=extract_zuru("current.xlsx",content)
    assert [r["fields"]["item"] for r in records]==["00123","00999"]
    record,notes=prepare_records(original(excel_bytes=content,item_no="00999"))
    assert record["fields"]["product"]=="CURRENT PRODUCT 00999"
    assert record["fields"]["invoice"]=="001100036670"
    assert any("已跳过 1" in note for note in notes)
    with pytest.raises(HTTPException): prepare_records(original(excel_bytes=content,item_no="unrelated"))
    record,notes=prepare_records(original(excel_bytes=content,item_no="internal",source_item="00999"))
    assert record["fields"]["item"]=="00999" and any("明确指定" in note for note in notes)
    with pytest.raises(HTTPException):prepare_records(original("zumiez",source_item="unrelated"))


def test_known_item_formula_error_blocks_generation():
    book=Workbook();sheet=book.active
    for row, values in enumerate((("Item No:", "00123"),("Product Description:","#N/A"),("Brand:","ZURU"),("PI No.","11")),1):
        for col,value in enumerate(values,1):sheet.cell(row,col,value)
    stream=io.BytesIO();book.save(stream);book.close()
    with pytest.raises(HTTPException) as error:extract_zuru("bad.xlsx",stream.getvalue())
    assert "公式错误" in error.value.detail


def test_ftc_profile_preserves_department_and_customer_article_without_a_generic_table():
    from app.services.carton_supplier_mark_pdf import prepare_pdf
    book=Workbook();sheet=book.active
    for cell,value in {"A1":"Farmers","A2":"Item:","B2":"00123","A3":"FTC","B3":"Dept #","C3":"LOI #","E3":"Order #",
        "A4":"FTC","B4":"/D685/","C4":"SSF248/","E4":"4504244551",
        "A5":"ARTICLE Number:","B5":"007044498","A6":"ARTICLE - Quantity:","B6":"8 PCS",
        "A7":"Carton number:","B7":"of 120","A8":"Description:","B8":"CURRENT PRODUCT",
        "A9":"Gross Weight:","B9":"2.800 KGS","A10":"Measurement:","B10":"29 x 21 x 31 CM",
        "A11":"Barcode Number:","B11":"00193052021172"}.items():sheet[cell]=value
    stream=io.BytesIO();book.save(stream);book.close()
    source=original(excel_bytes=stream.getvalue())
    content,_,_=prepare_pdf(source);text=PdfReader(io.BytesIO(content)).pages[0].extract_text()
    assert "New Zealand" in text and "007044498" in text and "4504244551" in text
    assert "Vendor name" not in text and "00193052021172" in text


def test_missing_zumiez_counts_are_explicit_and_order_or_user_values_never_replace_excel():
    record,notes=prepare_records(original("zumiez",item_no="factory-item"))
    assert "quantity" not in record["fields"] and any("缺少" in note for note in notes)
    record,notes=prepare_records(original("zumiez",supplements=dict(quantity="12 PCS",carton_total="42"),order_fields=dict(quantity="9")))
    assert record["fields"]["quantity"]=="12 PCS" and record["fields"]["customer_po"]=="00484860"
    with pytest.raises(HTTPException) as error:prepare_records(original(supplements=dict(gross_weight="5.76 KGS")))
    assert "不能覆盖原稿" in error.value.detail
    record,_=prepare_records(original(order_fields=dict(gross_weight="9 KGS")))
    assert record["fields"]["gross_weight"]=="5.93 KGS"


def test_esal_renderer_uses_fresh_text_and_requires_fixed_artwork_in_all_modes():
    from app.services.carton_supplier_mark_pdf import prepare_pdf
    for mode,pages in (("front_side",1),("front_only",1),("separate_pages",2)):
        source=original();source["layout"]["config"]["mode"]=mode
        content,count,_=prepare_pdf(source);reader=PdfReader(io.BytesIO(content))
        assert len(reader.pages)==count==pages
        if mode!="front_only":
            text="\n".join(page.extract_text() for page in reader.pages)
            assert "5.93 KGS" in text and "00193052021172" in text and "CURRENT PRODUCT" in text
    source=original();source["layout"]["config"]["front_artwork_region"]=None
    with pytest.raises(HTTPException) as error:prepare_pdf(source)
    assert "ESAL" in error.value.detail


def test_issued_outer_fields_convert_declared_inches_and_do_not_guess_ambiguous_or_unknown_units(monkeypatch):
    from app.services import carton_supplier_mark_pdf as service
    snapshot=dict(lines=[dict(packaging_type="外箱",specification="10 x 20 x 30",dimension_unit="in",
        usage_quantity="12",after_required_quantity="42",net_weight_kg="0",gross_weight_kg="5.04")])
    monkeypatch.setattr(service.portal,"latest_issue",lambda db,order:object())
    monkeypatch.setattr(service.portal,"_purchase_order_snapshot",lambda issue:snapshot)
    db=SimpleNamespace(get=lambda *args:object())
    assert service.issued_outer_fields(db,"order")==dict(dimensions="25.4 x 50.8 x 76.2 CM",quantity="12",carton_total="42",net_weight="0 KGS",gross_weight="5.04 KGS")
    snapshot["lines"][0]["dimension_unit"]=""
    assert "dimensions" not in service.issued_outer_fields(db,"order")
    snapshot["lines"].append(dict(snapshot["lines"][0]))
    assert service.issued_outer_fields(db,"order")=={}


def test_zumiez_api_generation_persists_current_supplements_with_provenance_and_no_qc_approval(monkeypatch):
    from test_molding_sample_api import make_client
    from test_carton_supplier_mark_pdf import prepare_layout,BASE
    from app.models.carton_mark import CartonMarkAsset,CartonMarkLayout,CartonMarkTemplate
    from app.models.carton_procurement import CartonAuditEvent
    import sqlalchemy as sa
    with make_client(monkeypatch) as client:
        _,payload,_=prepare_layout(client)
        from app.db import SessionLocal
        with SessionLocal() as db:
            asset=db.get(CartonMarkAsset,payload["excel_asset_id"]);content=excel("zumiez")
            asset.content=content;asset.size_bytes=len(content);asset.sha256=hashlib.sha256(content).hexdigest()
            layout=db.get(CartonMarkLayout,payload["layout_id"])
            layout.config_json=MarkLayoutConfig(series="zuru",logo_region=dict(x=10,y=10,width=20,height=20)).model_dump_json()
            db.commit()
        response=client.post(BASE,json={**payload,"supplements":dict(quantity="12 PCS",carton_total="42")})
        assert response.status_code==201,response.text
        result=response.json();assert result["asset"]["kind"]=="pdf"
        with SessionLocal() as db:
            event=db.scalar(sa.select(CartonAuditEvent).where(CartonAuditEvent.entity_id==result["asset"]["id"],CartonAuditEvent.event_type=="CARTON_MARK_ASSET_CREATED"))
            assert json.loads(event.detail_json)["conversion"]["supplements"]==dict(quantity="12 PCS",carton_total="42")
            content=bytes(db.get(CartonMarkAsset,result["asset"]["id"]).content)
            assert db.scalar(sa.select(sa.func.count()).select_from(CartonMarkTemplate))==0
        assert "12 PCS" in PdfReader(io.BytesIO(content)).pages[0].extract_text()
        assert client.post(BASE,json={**payload,"supplements":{"unit_price":"12"}}).status_code==422
