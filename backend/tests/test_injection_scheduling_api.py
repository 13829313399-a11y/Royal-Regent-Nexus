from io import BytesIO
import importlib
import sys
from datetime import datetime
from pathlib import Path
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import to_excel


TEST_TMP_DIR = Path(__file__).resolve().parents[1] / ".pytest-tmp"
BACKEND_DIR = Path(__file__).resolve().parents[1]
ADMIN_TEST_PASSWORD = "AdminSeed123!"

if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))


def make_client(monkeypatch) -> TestClient:
    TEST_TMP_DIR.mkdir(exist_ok=True)
    database_url = (
        f"sqlite:///{TEST_TMP_DIR / f'injection_scheduling_{uuid4().hex}.db'}"
    )
    monkeypatch.setenv("DATABASE_URL", database_url)
    monkeypatch.setenv("SEED_ADMIN_PASSWORD", ADMIN_TEST_PASSWORD)
    monkeypatch.setenv("AUTHZ_MODE", "enforce")
    monkeypatch.setenv("AUTHZ_WRITES_ENABLED", "false")

    for module_name in list(sys.modules):
        if module_name == "app" or module_name.startswith("app."):
            del sys.modules[module_name]

    main = importlib.import_module("app.main")
    return TestClient(main.app)


def ensure_user(
    username: str,
    role_id: str,
    *,
    factory_id: str,
    department: str = "molding",
) -> None:
    db_module = importlib.import_module("app.db")
    auth_models = importlib.import_module("app.models.auth")
    auth_service = importlib.import_module("app.services.auth")
    user_id = f"user-{username}"
    with db_module.SessionLocal() as db:
        salt, password_hash = auth_service.make_password_hash("123456")
        db.add(
            auth_models.AuthUser(
                id=user_id,
                username=username,
                display_name=username,
                password_salt=salt,
                password_hash=password_hash,
                status="active",
                force_password_change=0,
                created_at=auth_service.now_text(),
                updated_at=auth_service.now_text(),
            )
        )
        db.add(
            auth_models.AuthUserRole(
                id=f"{user_id}:{role_id}:{factory_id}:{department}",
                user_id=user_id,
                role_id=role_id,
                factory_id=factory_id,
                department=department,
            )
        )
        db.commit()


def login(client: TestClient, username: str) -> dict:
    response = client.post(
        "/api/auth/login",
        json={"username": username, "password": "123456"},
    )
    assert response.status_code == 200
    return response.json()


def build_huakang_workbook() -> bytes:
    workbook = Workbook()
    plan = workbook.active
    plan.title = "排期表"
    plan.cell(3, 2, "机台")
    plan.cell(3, 7, "模号")
    plan.cell(3, 9, "订单")
    plan.cell(3, 14, "剩余")
    plan.cell(4, 1, 1)
    plan.cell(4, 2, 1)
    plan.cell(4, 7, "60A 机台")
    plan.cell(5, 2, 1)
    plan.cell(5, 4, "▲ 特急订单")
    plan.cell(5, 5, 12)
    plan.cell(5, 5).number_format = "00000"
    plan.cell(5, 6, "60A")
    plan.cell(5, 7, "MOLD-001")
    plan.cell(5, 8, "测试产品")
    plan.cell(5, 9, 34)
    plan.cell(5, 9).number_format = "000000"
    plan.cell(5, 10, 7)
    plan.cell(5, 10).number_format = "000"
    plan.cell(5, 11, 2)
    plan.cell(5, 12, 1000)
    plan.cell(5, 13, 100)
    plan.cell(5, 14, 900)
    plan.cell(5, 15, 450)
    plan.cell(5, 16, "ABS")
    plan.cell(5, 17, "20%")
    plan.cell(5, 18, "黑色")
    plan.cell(5, 19, "SP-001")
    plan.cell(5, 20, 95)
    plan.cell(5, 21, 105)
    plan.cell(5, 22, 1.25)
    plan.cell(5, 23, 12.3)
    plan.cell(5, 24, to_excel(datetime(2026, 7, 20)))
    plan.cell(5, 25, to_excel(datetime(2026, 7, 28)))
    plan.cell(5, 26, to_excel(datetime(2026, 8, 1)))
    plan.cell(5, 27, 0.0625)
    plan.cell(5, 28, 0.0208333333)
    plan.cell(5, 29, 0.0833333333)
    plan.cell(5, 30, 0.0416666667)
    plan.cell(5, 31, to_excel(datetime(2026, 7, 28, 8)))
    plan.cell(5, 32, to_excel(datetime(2026, 7, 30, 8)))
    plan.cell(5, 33, "2026-07")
    plan.cell(5, 34, to_excel(datetime(2026, 7, 31)))
    plan.cell(5, 35, -1.5)
    plan.cell(5, 36, 2)
    plan.cell(5, 37, "否")
    plan.cell(5, 38, "20:00")
    plan.cell(5, 39, "12:00")
    plan.cell(5, 40, 225)
    plan.cell(5, 41, 5)
    plan.cell(5, 42, "=AO5+5")
    plan.cell(5, 43, 220)
    plan.cell(5, 44, 230)

    machine = workbook.create_sheet("厂区现有啤机")
    machine.cell(4, 1, 1)
    machine.cell(4, 2, "博创高速机")
    machine.cell(4, 3, "BS400")
    machine.cell(4, 4, "60A")
    machine.cell(4, 6, "1225G/43.2OZ")
    machine.cell(4, 7, "725MM*695MM")
    machine.cell(4, 8, 400)
    machine.cell(4, 9, "高速")
    machine.cell(4, 10, "五轴双臂")

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def build_huakang_workbook_with_repeated_order() -> bytes:
    workbook = load_workbook(BytesIO(build_huakang_workbook()))
    plan = workbook["排期表"]
    for column in range(1, 45):
        source = plan.cell(5, column)
        target = plan.cell(6, column, source.value)
        target.number_format = source.number_format
    plan.cell(6, 4, "后续同单任务")
    plan.cell(6, 10, 8)
    plan.cell(6, 10).number_format = "000"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def build_huaxing_workbook() -> bytes:
    workbook = Workbook()
    plan = workbook.active
    plan.title = "计划表"
    plan.cell(3, 2, "机台")
    plan.cell(3, 7, "模号")
    plan.cell(3, 9, "订单")
    plan.cell(3, 14, "剩余")
    plan.cell(4, 1, "旧1")
    plan.cell(4, 2, "旧1")
    plan.cell(4, 7, "50A 机台")
    plan.cell(5, 2, "旧1")
    plan.cell(5, 5, "▲")
    plan.cell(5, 6, "50A")
    plan.cell(5, 7, "HX-MOLD-001")
    plan.cell(5, 8, "华兴产品")
    plan.cell(5, 9, "HX-ORDER-001")
    plan.cell(5, 10, "HX-ITEM-001")
    plan.cell(5, 12, 500)
    plan.cell(5, 13, 20)
    plan.cell(5, 14, 480)
    plan.cell(5, 15, 240)
    plan.cell(5, 17, "透明")
    plan.cell(5, 19, "PC")
    plan.cell(5, 20, 80)
    plan.cell(5, 28, "2026-08-02")

    machine = workbook.create_sheet("机台完成时间")
    machine.cell(3, 1, "旧1")
    machine.cell(3, 3, "50A 400T")
    machine.cell(3, 4, "普通")
    machine.cell(3, 5, "双臂五轴")

    mold = workbook.create_sheet("机安")
    mold.cell(1, 1, "模号")
    mold.cell(2, 1, "HX-MOLD-001")
    mold.cell(2, 2, "华兴产品")
    mold.cell(2, 7, "50A")
    mold.cell(2, 8, "双臂")
    mold.cell(2, 12, "PC")
    mold.cell(2, 13, 80)

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_parser_recognizes_both_supported_templates_without_cross_factory_inference():
    parser = importlib.import_module(
        "app.services.injection_scheduling_import"
    ).parse_injection_schedule_workbook
    huakang = parser(
        build_huakang_workbook(),
        factory_id="huakang-b",
        source_file_name="huakang.xlsx",
        business_date="2026-07-28",
    )
    huaxing = parser(
        build_huaxing_workbook(),
        factory_id="huaxing",
        source_file_name="huaxing.xlsx",
        business_date="2026-07-28",
    )
    assert huakang.summary["machineCount"] == 1
    assert huakang.summary["taskCount"] == 1
    assert huakang.summary["canConfirm"] is True
    assert huaxing.summary["machineCount"] == 1
    assert huaxing.summary["taskCount"] == 1
    assert huaxing.summary["canConfirm"] is True
    assert huaxing.normalized["machines"][0]["completeness"] == "needs_review"
    huakang_order = huakang.normalized["orders"][0]
    worksheet = huakang_order["requirement"]["worksheet"]
    restricted = huakang_order["requirement"]["restrictedWorksheet"]
    assert huakang_order["itemNo"] == "00012"
    assert huakang_order["orderNo"] == "000034"
    assert worksheet["warehouse"] == "007"
    assert worksheet["setQuantity"] == 2
    assert worksheet["waterRatio"] == "20%"
    assert worksheet["colorPowder"] == "SP-001"
    assert worksheet["materialWeightKg"] == 1.25
    assert worksheet["deliveryDueAt"].startswith("2026-08-01")
    assert worksheet["moldChangeReferenceHours"] == 1.5
    assert worksheet["deliverySlackDays"] == -1.5
    assert worksheet["allocatedMaterialQuantity"] is None
    assert worksheet["sourceSheet"] == "排期表"
    assert worksheet["sourceRow"] == 5
    assert restricted["unitPricePerShot"] == 12.3
    assert any(
        issue.code == "formula_cache_missing"
        and issue.field == "allocated_material"
        and issue.source_row == 5
        for issue in huakang.issues
    )
    assert huakang.normalized["tasks"][0]["slackHours"] == -36

    repeated = parser(
        build_huakang_workbook_with_repeated_order(),
        factory_id="huakang-b",
        source_file_name="huakang-repeated.xlsx",
        business_date="2026-07-28",
    )
    assert repeated.summary["orderCount"] == 1
    assert repeated.summary["taskCount"] == 2
    assert [
        task["worksheet"]["warehouse"]
        for task in repeated.normalized["tasks"]
    ] == ["007", "008"]

    try:
        parser(
            build_huakang_workbook(),
            factory_id="huaxing",
            source_file_name="wrong-factory.xlsx",
            business_date="2026-07-28",
        )
    except ValueError as exc:
        assert "华康B模板" in str(exc)
    else:
        raise AssertionError("cross-factory workbook import must fail")


def test_api_import_version_publish_and_rollback_enforce_factory_and_role(monkeypatch):
    with make_client(monkeypatch) as client:
        huakang_api_workbook = build_huakang_workbook_with_repeated_order()
        anonymous = client.get(
            "/api/injection-scheduling/snapshots/current",
            params={"factory_id": "huakang-b"},
        )
        assert anonymous.status_code == 401

        ensure_user(
            "schedule-clerk",
            "position_molding_clerk",
            factory_id="huakang-b",
        )
        clerk = login(client, "schedule-clerk")
        assert "injection_scheduling:import" in clerk["permissions"]
        assert "injection_scheduling:publish" not in clerk["permissions"]

        preview = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huakang-b", "business_date": "2026-07-28"},
            files={
                "file": (
                    "huakang.xlsx",
                    huakang_api_workbook,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert preview.status_code == 200, preview.text
        preview_body = preview.json()
        assert preview_body["summary"]["taskCount"] == 2
        assert preview_body["summary"]["canConfirm"] is True
        before_confirm = client.get(
            "/api/injection-scheduling/snapshots/current",
            params={"factory_id": "huakang-b"},
        )
        assert before_confirm.status_code == 404

        wrong_factory = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huaxing", "business_date": "2026-07-28"},
            files={
                "file": (
                    "huakang.xlsx",
                    build_huakang_workbook(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert wrong_factory.status_code == 403

        confirm = client.post(
            f"/api/injection-scheduling/imports/{preview_body['batch_id']}/confirm",
            json={
                "factory_id": "huakang-b",
                "preview_revision": preview_body["preview_revision"],
                "reason": "确认测试排程",
            },
        )
        assert confirm.status_code == 200, confirm.text
        assert confirm.json()["snapshot"]["factoryId"] == "huakang-b"
        assert len(confirm.json()["snapshot"]["tasks"]) == 2
        task_worksheets = [
            task["worksheet"]
            for task in confirm.json()["snapshot"]["tasks"]
        ]
        assert [item["warehouse"] for item in task_worksheets] == ["007", "008"]
        task_worksheet = task_worksheets[0]
        assert task_worksheet["deliverySlackDays"] == -1.5
        assert task_worksheet["moldChangeReferenceHours"] == 1.5
        assert "unitPricePerShot" not in task_worksheet
        assert "restrictedWorksheet" not in confirm.text
        first_plan_id = confirm.json()["plan_id"]

        second_preview = client.post(
            "/api/injection-scheduling/imports/preview",
            data={"factory_id": "huakang-b", "business_date": "2026-07-28"},
            files={
                "file": (
                    "huakang.xlsx",
                    huakang_api_workbook,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert second_preview.status_code == 200
        assert second_preview.json()["preview_revision"] == 2
        second_confirm = client.post(
            f"/api/injection-scheduling/imports/{second_preview.json()['batch_id']}/confirm",
            json={
                "factory_id": "huakang-b",
                "preview_revision": 2,
                "reason": "验证同厂重新导入",
            },
        )
        assert second_confirm.status_code == 200, second_confirm.text
        assert second_confirm.json()["plan_id"] != first_plan_id

        current = client.get(
            "/api/injection-scheduling/snapshots/current",
            params={"factory_id": "huakang-b"},
        )
        assert current.status_code == 200
        assert current.json()["revision"] == 1

        forbidden_publish = client.post(
            "/api/injection-scheduling/plans/publish",
            json={
                "factory_id": "huakang-b",
                "revision": 1,
                "reason": "文员不可发布",
            },
        )
        assert forbidden_publish.status_code == 403

        client.post("/api/auth/logout")
        ensure_user(
            "schedule-supervisor",
            "position_molding_supervisor",
            factory_id="huakang-b",
        )
        supervisor = login(client, "schedule-supervisor")
        assert "injection_scheduling:publish" in supervisor["permissions"]
        assert "injection_scheduling:rollback" in supervisor["permissions"]

        conflict = client.post(
            "/api/injection-scheduling/plans/draft",
            json={
                "factory_id": "huakang-b",
                "revision": 999,
                "snapshot": current.json()["snapshot"],
                "reason": "验证版本冲突",
            },
        )
        assert conflict.status_code == 409

        publish = client.post(
            "/api/injection-scheduling/plans/publish",
            json={
                "factory_id": "huakang-b",
                "revision": 1,
                "reason": "主管发布验收",
            },
        )
        assert publish.status_code == 200, publish.text
        version = publish.json()["version"]

        published = client.get(
            "/api/injection-scheduling/published/current",
            params={"factory_id": "huakang-b"},
        )
        assert published.status_code == 200
        assert published.json()["version"] == version
        assert published.json()["snapshot"]["factoryId"] == "huakang-b"
        assert [
            task["worksheet"]["warehouse"]
            for task in published.json()["snapshot"]["tasks"]
        ] == ["007", "008"]
        assert "unitPricePerShot" not in published.json()["snapshot"]["tasks"][0]["worksheet"]

        current_after_publish = client.get(
            "/api/injection-scheduling/snapshots/current",
            params={"factory_id": "huakang-b"},
        )
        assert current_after_publish.status_code == 200
        assert current_after_publish.json()["snapshot"]["plan"]["status"] == "published"
        assert (
            current_after_publish.json()["snapshot"]["plan"]["publishedAt"]
            == publish.json()["published_at"]
        )

        rollback = client.post(
            "/api/injection-scheduling/plans/rollback",
            json={
                "factory_id": "huakang-b",
                "version": version,
                "revision": 1,
                "reason": "主管回滚验收",
            },
        )
        assert rollback.status_code == 200, rollback.text
        assert rollback.json()["revision"] == 2
        assert rollback.json()["snapshot"]["plan"]["status"] == "draft"
        assert "publishedAt" not in rollback.json()["snapshot"]["plan"]
        assert [
            task["worksheet"]["warehouse"]
            for task in rollback.json()["snapshot"]["tasks"]
        ] == ["007", "008"]
