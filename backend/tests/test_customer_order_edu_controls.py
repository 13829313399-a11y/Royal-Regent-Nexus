"""Synthetic EDU products and quantities cover reservation and manual controls."""
from io import BytesIO

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from app.services import customer_order_unified as service
from app.services.customer_order_manual import decorate_manual_resolution_policy
from app.services.customer_order_edu import allocate_references, prepare_controls
from test_customer_order_huaxing_unified import template, data
from test_edu_order_identity import po_bytes as _po_bytes, preview_args


PRODUCT = 'F12-ZZ901-C901'


def po_bytes(contracts, *, quantities=None, pos=None):
    return _po_bytes(
        contracts,
        products=[PRODUCT + '0000'] * len(contracts),
        quantities=quantities if quantities is not None else [600] * len(contracts),
        pos=pos,
    )


def stock_book():
    wb = template('edu')
    sheet = wb['ITEM表']
    for r, po, ref in [(4, 'P1', 'EDUHX00101'), (5, 'P2', 'EDUHX00103')]:
        for c, value in {4: 'PO订单', 5: ref, 6: po, 8: PRODUCT + '0000', 11: 0, 12: '1/16'}.items():
            sheet.cell(r, c, value)
        for c in range(1, 29):
            sheet.cell(r, c).fill = PatternFill('solid', fgColor='4EA72E')
    return wb


@pytest.mark.parametrize('po,expected', [('PO#P1', '4'), ('P2', '5'), ('NEW', '5')])
def test_po_match_then_last_green_row_and_normalized_pack(po, expected):
    args = preview_args(po_bytes(['C1'], pos=[po]), stock_book())
    preview = service.create_unified_customer_preview(**args)
    row = preview['rows'][0]
    assert row['edu_reservation_row'] == expected
    assert row['reference_no'] == 'EDUHX00104'
    assert row['product_no'] == PRODUCT
    assert row['source_product_no'] == PRODUCT + '0000'
    assert row['source_order_date'] == '2026-09-29'
    assert row['units_per_carton'] == '16'
    assert not any(i['severity'] == 'blocked' for i in row['issues'])
    output, _, _ = service.export_unified_customer_schedule(**args)
    book = openpyxl.load_workbook(BytesIO(output))
    sheet = book['ITEM表']
    reserve_ref = 'EDUHX00101' if expected == '4' else 'EDUHX00103'
    reserve = next(r for r in range(4, sheet.max_row + 1) if sheet.cell(r, 5).value == reserve_ref)
    assert sheet.cell(reserve, 11).value == '=0-600'
    assert book['PO订单']['B2'].value == 'EDUHX00104'
    assert str(book['PO订单']['A2'].value).startswith('2026-09-29')


def test_allocator_considers_other_sheets_and_child_numbers():
    wb = template('edu')
    wb.create_sheet('PO订单')['B2'] = 'EDUHX00990'
    wb['正单评审表']['D4'] = 'EDUHX01000-2'
    rows = [dict(contract_no='C', po_no='P', product_no=p) for p in ['I1', 'I2', 'I1']]
    allocate_references(wb, rows)
    assert [r['reference_no'] for r in rows] == ['EDUHX01001', 'EDUHX01002', 'EDUHX01001']


def export_with_override(args, field, value):
    preview = service.create_unified_customer_preview(**args)
    decorate_manual_resolution_policy(preview)
    row = preview['rows'][0]
    issue = next(i for i in row['issues'] if i.get('edit_field') == field)
    override = dict(row_id=row['id'], issue_key=issue['skip_key'], field=field, value=value)
    return service.export_unified_customer_schedule(**args, manual_overrides=[override], skipped_issue_keys={issue['skip_key']})


@pytest.mark.parametrize('value', ['0', '-1', 'NaN', 'garbage'])
def test_invalid_manual_pack_cannot_create_false_cartons(value):
    args = preview_args(po_bytes(['C1']), stock_book())
    with pytest.raises(service.CustomerOrderUnifiedError, match='装箱数'):
        export_with_override(args, 'units_per_carton', value)


def test_manual_pack_is_used_in_excel_and_resolved_preview():
    args = preview_args(po_bytes(['C1']), stock_book())
    output, _, preview = export_with_override(args, 'units_per_carton', '24')
    row = preview['rows'][0]
    assert row['carton_count'] == '25'
    book = openpyxl.load_workbook(BytesIO(output))
    sheet = book['ITEM表']
    r = next(r for r in range(4, sheet.max_row + 1) if sheet.cell(r, 5).value == 'EDUHX00104')
    assert sheet.cell(r, 12).value == 24
    assert f'K{r}/L{r}' in sheet.cell(r, 13).value


def test_wrong_manual_stock_row_rejected():
    with pytest.raises(service.CustomerOrderUnifiedError, match='备料行号'):
        export_with_override(preview_args(po_bytes(['C1']), stock_book()), 'edu_reservation_row', '7')


def test_repeat_prefers_its_own_manual_pack_over_other_orders_or_stock():
    args = preview_args(po_bytes(['C1']), stock_book())
    output, _, _ = export_with_override(args, 'units_per_carton', '24')
    retry = service.create_unified_customer_preview(**{**args, 'schedule_content': output})
    row = retry['rows'][0]
    assert row['units_per_carton'] == '24'
    assert row['carton_count'] == '25'
    assert not any(i.get('can_edit') for i in row['issues'])
    with pytest.raises(service.CustomerOrderUnifiedError, match='重复单'):
        service.export_unified_customer_schedule(**{**args, 'schedule_content': output}, manual_overrides=[
            dict(row_id=row['id'], field='units_per_carton', value='16', issue_key='forged')])


def test_repeated_confirmed_order_never_reduces_reservation_twice():
    args = preview_args(po_bytes(['C1']), stock_book())
    output, _, _ = service.export_unified_customer_schedule(**args)
    retry = {**args, 'schedule_content': output}
    preview = service.create_unified_customer_preview(**retry)
    keys = {i['skip_key'] for r in preview['rows'] for i in r['issues'] if i.get('can_skip')}
    second, _, resolved = service.export_unified_customer_schedule(**retry, skipped_issue_keys=keys)
    before, after = [openpyxl.load_workbook(BytesIO(content)) for content in (output, second)]
    assert resolved['rows'][0]['reference_no'] == 'EDUHX00104'
    for name in before.sheetnames:
        assert {(r,c): v.value for (r,c),v in before[name]._cells.items() if v.value is not None} == {
            (r,c): v.value for (r,c),v in after[name]._cells.items() if v.value is not None}


def test_quantity_conflict_cannot_be_turned_into_optional_confirmation():
    args = preview_args(po_bytes(['C1']), stock_book())
    output, _, _ = service.export_unified_customer_schedule(**args)
    args.update(schedule_content=output, po_files=[('changed.xls', po_bytes(['C1'], quantities=[576]))])
    preview = decorate_manual_resolution_policy(service.create_unified_customer_preview(**args))
    issue = next(i for i in preview['rows'][0]['issues'] if i['code'] == 'existing_quantity_conflict')
    assert not issue['can_skip'] and not issue['can_edit']
    with pytest.raises(service.CustomerOrderUnifiedError):
        service.export_unified_customer_schedule(**args, skipped_issue_keys={issue['skip_key']})
