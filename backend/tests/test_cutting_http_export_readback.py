"""Load evidence accepts XLSX/ZIP and rejects missing or damaged volumes."""
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile, ZIP_STORED
from hashlib import sha256
import xml.etree.ElementTree as ET
import json

import httpx
import pytest
from openpyxl import load_workbook

from app.services import cutting_export_workbook as workbook
from test_cutting_http_capacity import validate_workbook


def make_response(monkeypatch, divided):
    report = dict(kind='period',month='2026-10',token='fixed',captured_at='2026-10-21',metadata=[],sheets=[
        dict(name='费用汇总',headers=['币种','本期正常费用','期末累计完工费用','期末在制费用','本期损耗费用'],
             rows=[['HKD','8.00','4.80','3.20','2.00']]),
        dict(name='明细',headers=['行号'],rows=[[1],[2]])])
    if divided:
        monkeypatch.setattr(workbook,'plan',lambda report:[
            [dict(table=0,start=0,end=1)], [dict(table=1,start=0,end=2)]])
    content,extension = workbook.package(report,SimpleNamespace(id='synthetic',display_name='验证员'),
        lambda value,actor:workbook.render_workbook(value,actor,'分卷验收'))
    response = httpx.Response(200,content=content,headers={'cache-control':'no-store',
        'content-type':'application/zip' if extension=='zip' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'x-cutting-report-token':'fixed'})
    return response,workbook.delivery(report)


@pytest.mark.parametrize('divided',[False,True])
@pytest.mark.parametrize('full',[False,True])
def test_readback_accepts_single_and_divided_reports(monkeypatch,divided,full):
    response,delivery = make_response(monkeypatch,divided)
    validate_workbook(response,expected_delivery=delivery,full_readback=full)


@pytest.mark.parametrize('damage',['missing_file','digest','coverage','token','physical_rows'])
def test_readback_rejects_incomplete_or_inconsistent_archive(monkeypatch,damage):
    response,delivery = make_response(monkeypatch,True)
    with ZipFile(BytesIO(response.content)) as original:
        files = {name:original.read(name) for name in original.namelist()}
    manifest = json.loads(files['manifest.json'])
    if damage=='missing_file':
        files.pop(manifest['files'][-1]['name'])
    elif damage=='digest':
        manifest['files'][0]['sha256']='invalid'
    elif damage=='coverage':
        manifest['coverage'][1]['first']=2
    elif damage=='token':
        manifest['token']='another'
    else:
        # Even a self-consistent manifest must match the real worksheet size.
        manifest['coverage'][1].update(last=3,rows=3,total=3)
        delivery['manifest']=manifest['coverage']
    files['manifest.json']=json.dumps(manifest).encode()
    output = BytesIO()
    with ZipFile(output,'w',ZIP_STORED) as damaged:
        for name,content in files.items():
            damaged.writestr(name,content)
    bad = httpx.Response(200,content=output.getvalue(),headers=response.headers)
    with pytest.raises(AssertionError):
        validate_workbook(bad,expected_delivery=delivery)


@pytest.mark.parametrize('damage',['embedded_manifest','phantom_row'])
def test_final_readback_checks_inner_manifest_and_actual_stream(monkeypatch,damage):
    response,delivery = make_response(monkeypatch,True)
    with ZipFile(BytesIO(response.content)) as original:
        files = {name:original.read(name) for name in original.namelist()}
    manifest = json.loads(files['manifest.json'])
    item = manifest['files'][1]
    book = load_workbook(BytesIO(files[item['name']]))
    if damage == 'embedded_manifest':
        book['完整性清单']['D4'] = 99
        output = BytesIO()
        book.save(output)
        raw = output.getvalue()
    else:
        sheet_path = f"xl/worksheets/sheet{book.sheetnames.index('明细')+1}.xml"
        with ZipFile(BytesIO(files[item['name']])) as original:
            parts = {name:original.read(name) for name in original.namelist()}
        root = ET.fromstring(parts[sheet_path])
        data = root.find('{*}sheetData')
        data.remove(next(row for row in data if row.attrib['r']=='4'))
        # Keep the declared dimension intact: max_row alone would still pass.
        parts[sheet_path] = ET.tostring(root)
        output = BytesIO()
        with ZipFile(output,'w',ZIP_STORED) as changed:
            for name,content in parts.items():
                changed.writestr(name,content)
        raw = output.getvalue()
    book.close()
    files[item['name']] = raw
    item.update(bytes=len(raw),sha256=sha256(raw).hexdigest())
    files['manifest.json'] = json.dumps(manifest).encode()
    output = BytesIO()
    with ZipFile(output,'w',ZIP_STORED) as changed:
        for name,content in files.items():
            changed.writestr(name,content)
    bad = httpx.Response(200,content=output.getvalue(),headers=response.headers)
    validate_workbook(bad,expected_delivery=delivery)
    with pytest.raises(AssertionError):
        validate_workbook(bad,expected_delivery=delivery,full_readback=True)
