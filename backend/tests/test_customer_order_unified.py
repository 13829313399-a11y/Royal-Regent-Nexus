from __future__ import annotations

from io import BytesIO

import openpyxl

from app.services import customer_order_unified as unified


def _template_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for sheet_name, headers in (
        (unified.ORDER_SHEET, unified.ORDER_HEADERS),
        (unified.REVIEW_SHEET, unified.ORDER_HEADERS),
        (unified.ITEM_SHEET, unified.ITEM_HEADERS),
    ):
        worksheet = workbook.create_sheet(sheet_name)
        for column, header in enumerate(headers, start=1):
            worksheet.cell(3, column).value = header
        for row in range(4, 12):
            if sheet_name == unified.ITEM_SHEET:
                worksheet.cell(row, 13).value = f"=K{row}/L{row}"
            else:
                for column in range(1, 21):
                    if column not in {3, 4, 5, 6, 7, 14}:
                        worksheet.cell(row, column).value = f'=IF(D{row}="","","公式")'
        worksheet.cell(12, 1).value = "取消单"
        worksheet.cell(13, 2).value = "用户维护内容"
        if sheet_name != unified.ORDER_SHEET:
            worksheet.cell(20, 1).value = "已走货订单"

    item = workbook[unified.ITEM_SHEET]
    order = workbook[unified.ORDER_SHEET]
    review = workbook[unified.REVIEW_SHEET]
    for row in range(4, 7):
        item.cell(row, 4).value = f"CTR-DEMO-{row}"
        item.cell(row, 5).value = f"SO-DEMO-{row}"
        item.cell(row, 6).value = f"PO-DEMO-{row}"
        item.cell(row, 7).value = "演示客户"
        item.cell(row, 8).value = f"ITEM-DEMO-{row}"
        item.cell(row, 29).value = "用户示例日期码"
        item.cell(row, 30).value = "模拟数据"
        for summary in (order, review):
            summary.cell(row, 3).value = f"CTR-DEMO-{row}"
            summary.cell(row, 4).value = f"SO-DEMO-{row}"
            summary.cell(row, 5).value = f"PO-DEMO-{row}"
            summary.cell(row, 6).value = "演示客户"
            summary.cell(row, 7).value = f"ITEM-DEMO-{row}"
        order.cell(row, 14).value = 1

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _parsed_row(index: int = 1) -> dict:
    return {
        "id": f"buzzbee-{index}",
        "status": "valid",
        "status_label": "可导出",
        "row_role": "detail",
        "parent_product_no": "",
        "received_date": "2026-09-04",
        "po_no": f"PO-{index}",
        "contract_no": f"CTR-{index}",
        "reference_no": f"SO-{index}",
        "customer_country": "BuzzBee / US",
        "customer_name": "BuzzBee",
        "country": "US",
        "product_no": f"ITEM-{index}",
        "product_name_zh": f"产品{index}",
        "product_name_en": f"Product {index}",
        "quantity": "240",
        "units_per_carton": "24",
        "carton_count": "10",
        "standard": "ASTM",
        "unit_price_hkd": "12.5",
        "amount_hkd": "3000",
        "packaging": "彩盒",
        "line_q": "2026-09-20",
        "customer_q": "2026-09-21",
        "requested_ship_date": "2026-09-30",
        "input_template": "TEST_PO_V1",
        "target_template": "legacy",
        "item_sheet_name": "legacy",
        "source_po_file_name": "order.xlsx",
        "lineage": {},
        "issues": [],
        "date_code": "DC2609",
    }


def test_unified_template_validation_and_demo_history_ignored():
    content = _template_bytes()
    unified.ensure_unified_schedule("河源业务统一排期.xlsx", content)
    assert unified.read_unified_history(content) == []


def test_optional_common_item_fields_stay_blank_without_blocking(monkeypatch):
    row = _parsed_row()
    for field in (
        "order_type",
        "production_no",
        "product_name_zh",
        "product_name_en",
        "quantity",
        "units_per_carton",
        "standard",
        "packaging",
        "date_code",
        "system_status",
        "requested_ship_date",
    ):
        row[field] = ""
    row["issues"] = [
        {
            "severity": "blocked",
            "code": "missing_required_field",
            "field": "quantity",
            "message": "旧模板要求数量",
            "can_skip": False,
            "skip_key": "",
            "skip_label": "",
        }
    ]
    monkeypatch.setattr(unified, "_create_customer_rows", lambda *_args: ([row], [], "TEST_PO_V1"))

    output, _file_name, preview = unified.export_unified_customer_schedule(
        customer_code="buzzbee",
        factory_id="huaxing",
        received_date="2026-09-04",
        po_files=[("order.xlsx", b"po")],
        schedule_file_name="河源业务统一排期.xlsx",
        schedule_content=_template_bytes(),
    )

    assert preview["summary"]["blocked"] == 0
    workbook = openpyxl.load_workbook(BytesIO(output), data_only=False)
    try:
        item = workbook[unified.ITEM_SHEET]
        assert item["B7"].value is None
        assert item["I7"].value is None
        assert item["K7"].value is None
        assert item["L7"].value is None
        assert item["M7"].value is None
        assert item["N7"].value is None
        assert item["T7"].value is None
        assert item["V7"].value is None
        assert item["Z7"].value is None
    finally:
        workbook.close()


def test_export_only_uses_rows_above_cancel_and_leaves_ac_ad_for_user(monkeypatch):
    content = _template_bytes()
    monkeypatch.setattr(
        unified,
        "_create_customer_rows",
        lambda *_args: ([_parsed_row()], [], "TEST_PO_V1"),
    )

    output, file_name, preview = unified.export_unified_customer_schedule(
        customer_code="buzzbee",
        factory_id="huaxing",
        received_date="2026-09-04",
        po_files=[("order.xlsx", b"po")],
        schedule_file_name="河源业务统一排期.xlsx",
        schedule_content=content,
    )

    assert file_name.endswith("_BuzzBee新单.xlsx")
    assert preview["target_template"] == unified.huaxing.TARGET_TEMPLATE
    workbook = openpyxl.load_workbook(BytesIO(output), data_only=False)
    try:
        item = workbook[unified.ITEM_SHEET]
        order = workbook[unified.ORDER_SHEET]
        review = workbook[unified.REVIEW_SHEET]
        assert [order.cell(7, column).value for column in range(3, 8)] == [
            item.cell(7, column).value for column in range(4, 9)
        ]
        assert [review.cell(7, column).value for column in range(3, 8)] == [
            item.cell(7, column).value for column in range(4, 9)
        ]
        assert item["A7"].value.date().isoformat() == "2026-09-04"
        assert item["M7"].value == "=K7/L7"
        assert item["AC7"].value is None
        assert item["AD7"].value is None
        assert order["N7"].value is None
        assert item["E4"].value == "SO-DEMO-4"
        assert item["AD4"].value == "模拟数据"
        assert order["A12"].value == "取消单"
        assert review["A12"].value == "取消单"
        assert item["A12"].value == "取消单"
        assert order["B13"].value == "用户维护内容"
        assert review["A20"].value == "已走货订单"
    finally:
        workbook.close()


def test_export_inserts_before_cancel_without_changing_user_sections(monkeypatch):
    workbook = openpyxl.load_workbook(BytesIO(_template_bytes()), data_only=False)
    for sheet_name in unified.SHEETS:
        workbook[sheet_name].delete_rows(7, 5)
    compact = BytesIO()
    workbook.save(compact)
    workbook.close()
    monkeypatch.setattr(
        unified,
        "_create_customer_rows",
        lambda *_args: ([_parsed_row()], [], "TEST_PO_V1"),
    )
    output, _file_name, _preview = unified.export_unified_customer_schedule(
        customer_code="buzzbee",
        factory_id="huaxing",
        received_date="2026-09-04",
        po_files=[("order.xlsx", b"po")],
        schedule_file_name="河源业务统一排期.xlsx",
        schedule_content=compact.getvalue(),
    )
    workbook = openpyxl.load_workbook(BytesIO(output), data_only=False)
    try:
        for sheet_name in unified.SHEETS:
            worksheet = workbook[sheet_name]
            assert worksheet["A8"].value == "取消单"
            assert worksheet["B9"].value == "用户维护内容"
        assert workbook[unified.REVIEW_SHEET]["A16"].value == "已走货订单"
        assert workbook[unified.ORDER_SHEET]["H7"].value == '=IF(LEN(\'ITEM表\'!I7)=0,"",\'ITEM表\'!I7)'
        assert workbook[unified.REVIEW_SHEET]["H7"].value == '=IF(LEN(\'ITEM表\'!I7)=0,"",\'ITEM表\'!I7)'
        assert workbook[unified.ITEM_SHEET]["M7"].value == "=K7/L7"
        assert workbook[unified.ITEM_SHEET]["AC7"].value is None
        assert workbook[unified.ITEM_SHEET]["AD7"].value is None
    finally:
        workbook.close()


def test_export_appends_after_last_real_order_instead_of_filling_an_earlier_gap(monkeypatch):
    workbook = openpyxl.load_workbook(BytesIO(_template_bytes()), data_only=False)
    for sheet_name in unified.SHEETS:
        worksheet = workbook[sheet_name]
        for row in range(4, 7):
            for column in range(1, worksheet.max_column + 1):
                cell = worksheet.cell(row, column)
                if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    cell.value = None
    for row in (4, 6):
        item = workbook[unified.ITEM_SHEET]
        item.cell(row, 4).value = f"EXISTING-CTR-{row}"
        item.cell(row, 5).value = f"EXISTING-SO-{row}"
        item.cell(row, 8).value = f"EXISTING-ITEM-{row}"
        for sheet_name in (unified.ORDER_SHEET, unified.REVIEW_SHEET):
            worksheet = workbook[sheet_name]
            worksheet.cell(row, 3).value = f"EXISTING-CTR-{row}"
            worksheet.cell(row, 4).value = f"EXISTING-SO-{row}"
            worksheet.cell(row, 7).value = f"EXISTING-ITEM-{row}"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    monkeypatch.setattr(
        unified,
        "_create_customer_rows",
        lambda *_args: ([_parsed_row()], [], "TEST_PO_V1"),
    )
    output, _file_name, _preview = unified.export_unified_customer_schedule(
        customer_code="buzzbee",
        factory_id="huaxing",
        received_date="2026-09-04",
        po_files=[("order.xlsx", b"po")],
        schedule_file_name="河源业务统一排期.xlsx",
        schedule_content=buffer.getvalue(),
    )
    result = openpyxl.load_workbook(BytesIO(output), data_only=False)
    try:
        item = result[unified.ITEM_SHEET]
        assert item["E5"].value is None
        assert item["E6"].value == "EXISTING-SO-6"
        assert item["E7"].value == "SO-1"
        assert item["A12"].value == "取消单"
    finally:
        result.close()
