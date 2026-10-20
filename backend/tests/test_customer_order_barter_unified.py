"""Synthetic Barter orders cover duplicate handling and workbook preservation."""
from copy import copy
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
import pytest

from app.services import customer_order_unified as service
from app.services.customer_order_buzzbee import _encrypt_schedule
from app.services.huaxing_order_legacy import multi_po_parser as parser


CONTRACTS = ("90701", "90702", "90703", "90704")
QUANTITIES = (100, 150, 200, 250)
CUSTOMER_POS = ("0100000001", "0100000002", "0100000003", "0100000004")


def scan_pages():
    pages = []
    for contract, quantity, customer_po, day in zip(CONTRACTS, QUANTITIES, CUSTOMER_POS, (12, 14, 14, 23)):
        pages.append(f"""BARTER PURCHASE CONTRACT CONTRACT #: TNT/F{contract}
DELIVERY :OCT {day}, 2026
CUSTOMER :WAL-MART USA ({customer_po})
88111 SAMPLE TOY {quantity} PCS _12.3450/PCS - {quantity * 12.345:,.2f}
({quantity // 5} CTN)
5 PC PER MASTER
PLEASE APPLY DATE CODE WM88111RR-!?26TO BOTH ON PRODUCT & PACKAGING.
""")
    return pages


def template(*, missing_date_code=False, conflicting_date_code=False, quantity_conflict=False):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in ("ITEM表", "接单表", "正单评审表"):
        sheet = workbook.create_sheet(name)
        headers = service.ITEM_HEADERS if name == "ITEM表" else service.ORDER_HEADERS
        for col, value in enumerate(headers, 1):
            sheet.cell(3, col, value)
        sheet.cell(20, 1, "取消单")
        sheet.cell(21, 1, "历史取消记录")
    item = workbook["ITEM表"]
    for number, (contract, quantity, customer_po) in enumerate(zip(CONTRACTS, QUANTITIES, CUSTOMER_POS), 4):
        values = {4: f"TNT/F{contract}", 5: f"TNT/F{contract}", 6: customer_po,
                  7: "WM USA", 8: "88111", 9: "样例玩具", 10: "SAMPLE TOY",
                  11: quantity, 12: 5, 13: f"=K{number}/L{number}", 21: "WM88111RR-1026"}
        for column, value in values.items():
            cell = item.cell(number, column, value)
            cell.font = Font(name="Arial", size=11, color="FF245678")
            cell.fill = PatternFill("solid", fgColor="FFFFF2CC")
        item.row_dimensions[number].height = 27
        for name in ("接单表", "正单评审表"):
            sheet = workbook[name]
            for target_col, source_col in zip(range(3, 8), range(4, 9)):
                sheet.cell(number, target_col, item.cell(number, source_col).value)
            sheet.cell(number, 10, f"='ITEM表'!K{number}")
            sheet.cell(number, 14, 12.345)
    if missing_date_code:
        item["U4"] = None
    if conflicting_date_code:
        for cell in list(item[4]):
            item.cell(8, cell.column, cell.value)
        item["U8"] = "WM88111RR-1126"
    if quantity_conflict:
        item["K4"] = 105
    extra = workbook.create_sheet("已出发票")
    extra["A1"], extra["B1"] = "原有发票信息", "=2+2"
    extra["A1"].font = Font(name="Arial", bold=True, color="FF123456")
    extra.merge_cells("A2:D2")
    return workbook


def content(workbook):
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def arguments(workbook, encrypted=False):
    payload = content(workbook)
    return dict(customer_code="barter", factory_id="huaxing", received_date="2026-10-15",
                po_files=[("TNT90701-90704 Contract.pdf", b"%PDF-1.4 scan fixture")],
                schedule_file_name="百德统一排期.xlsx",
                schedule_content=_encrypt_schedule(payload) if encrypted else payload)


@pytest.fixture(autouse=True)
def observed_scan(monkeypatch):
    monkeypatch.setattr(parser, "read_pdf_pages", lambda *args, **kwargs: (scan_pages(), True))


@pytest.mark.parametrize("encrypted", [False, True])
def test_plain_and_encrypted_unified_sources_export_red_duplicates_preserving_history(encrypted):
    original = template()
    args = arguments(original, encrypted)
    preview = service.create_unified_customer_preview(**args)
    assert preview["summary"] == {"total": 4, "valid": 0, "warning": 4, "blocked": 0}
    for row in preview["rows"]:
        assert row["date_code"] == "WM88111RR-1026"
        assert "同合同" in row["lineage"]["date_code"]
        assert not any(issue["code"] == "missing_date_code" for issue in row["issues"])
        assert any(issue["code"] == "duplicate_order_warning" for issue in row["issues"])
    output, _, _ = service.export_unified_customer_schedule(**args)
    rendered = load_workbook(BytesIO(output), data_only=False)
    try:
        assert rendered.sheetnames == original.sheetnames
        for name in original.sheetnames:
            for row in original[name].iter_rows(max_row=7 if name != "已出发票" else 2):
                for cell in row:
                    generated = rendered[name][cell.coordinate]
                    assert generated.value == cell.value
                    for attribute in ("font", "fill", "border", "alignment", "protection", "number_format"):
                        assert copy(getattr(generated, attribute)) == copy(getattr(cell, attribute))
        for name in ("ITEM表", "接单表", "正单评审表"):
            sheet = rendered[name]
            identity_column = 5 if name == "ITEM表" else 4
            added = [r for r in range(8, sheet.max_row + 1)
                     if sheet.cell(r, identity_column).value in {f"TNT/F{n}" for n in CONTRACTS}]
            assert len(added) == 4
            for number in added:
                assert all(c.font.color.type == "rgb" and c.font.color.rgb == "FFFF0000" for c in sheet[number])
                if name == "ITEM表":
                    assert sheet.cell(number, 21).value == "WM88111RR-1026"
                else:
                    assert sheet.cell(number, 10).data_type == "f"
                    assert "'ITEM表'!K" in sheet.cell(number, 10).value
        assert str(rendered["已出发票"].merged_cells) == str(original["已出发票"].merged_cells)
    finally:
        rendered.close()
        original.close()


@pytest.mark.parametrize("missing,conflicting", [(True, False), (False, True)])
def test_missing_or_conflicting_same_order_code_stays_editable_without_using_another_order(missing, conflicting):
    workbook = template(missing_date_code=missing, conflicting_date_code=conflicting)
    args = arguments(workbook)
    preview = service.create_unified_customer_preview(**args)
    row = preview["rows"][0]
    assert row.get("date_code", "") == ""
    issue = next(issue for issue in row["issues"] if issue["code"] == "missing_date_code")
    assert issue["severity"] == "warning" and issue["can_edit"]
    assert issue["edit_field"] == "date_code"
    output, _, _ = service.export_unified_customer_schedule(**args)
    rendered = load_workbook(BytesIO(output))
    assert rendered["ITEM表"]["U9" if conflicting else "U8"].value is None
    rendered.close()
    output, _, _ = service.export_unified_customer_schedule(**args, manual_overrides=[
        {"row_id": row["id"], "field": "date_code", "value": "WM88111RR-1026"},
    ])
    rendered = load_workbook(BytesIO(output))
    assert rendered["ITEM表"]["U9" if conflicting else "U8"].value == "WM88111RR-1026"
    rendered.close()
    workbook.close()


def test_existing_quantity_conflict_still_blocks_export():
    workbook = template(quantity_conflict=True)
    args = arguments(workbook)
    preview = service.create_unified_customer_preview(**args)
    assert preview["rows"][0]["status"] == "blocked"
    assert any(i["code"] == "existing_quantity_conflict" for i in preview["rows"][0]["issues"])
    with pytest.raises(service.CustomerOrderUnifiedError, match="数量"):
        service.export_unified_customer_schedule(**args)
    workbook.close()
