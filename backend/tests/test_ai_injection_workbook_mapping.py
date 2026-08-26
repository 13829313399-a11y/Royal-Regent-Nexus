from __future__ import annotations

import asyncio
import hashlib
import json
from datetime import date
from io import BytesIO

import pytest
from app.core.config import Settings
from app.schemas.ai.workbook import (
    AIInjectionWorkbookMappingV1,
    AIModelInjectionWorkbookMappingV1,
)
from app.services.ai.injection_workbook_mapping_recognition import (
    recognize_injection_workbook_mapping,
    validate_mapping_sources,
)
from app.services.ai.providers import ProviderResponse, ProviderToolCall
from app.services.ai.workbook_recognition_packet import (
    build_workbook_recognition_packet,
    structural_layout_signature,
)
from openpyxl import Workbook


def _workbook(*, month: int = 1, shift_days: int = 2) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = f"{month}月"
    sheet["A1"] = f"2026年{month}月生产计划"
    sheet["A2"] = "审核编号"
    sheet["B2"] = f"PLAN-{month:02d}-001"
    headers = {
        "A3": "机台",
        "B3": "机台编号",
        "G3": "模具编号",
        "I3": "订单号",
        "L3": "订单数量",
        "M3": "完成数量",
        "AG3": "计划开始",
        "AH3": "计划完成",
    }
    for ref, value in headers.items():
        sheet[ref] = value
    start_column = 47  # AU
    for day_no in range(shift_days):
        left = start_column + day_no * 2
        right = left + 1
        sheet.merge_cells(start_row=2, start_column=left, end_row=2, end_column=right)
        sheet.cell(2, left, date(2026, month, day_no + 1))
        sheet.cell(3, left, "白班")
        sheet.cell(3, right, "夜班")
    sheet["A4"] = "16#"
    sheet["B4"] = "16#"
    sheet["B5"] = "16#"
    sheet["G5"] = "M-001"
    sheet["I5"] = "SO-001"
    sheet["L5"] = 1000
    sheet["M5"] = 120
    sheet["AG5"] = "2026-01-01 08:00"
    sheet["AH5"] = "2026-01-02 20:00"
    sheet["A8"] = "合计"
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def _mapping(
    source_hash: str, *, header_cell: str = "G3"
) -> AIInjectionWorkbookMappingV1:
    def field(name: str, column: str, transformer: str, cell: str | None = None):
        return {
            "canonical_field": name,
            "source_column": column,
            "header_cell": cell or f"{column}3",
            "transformer": transformer,
            "confidence": 0.96,
            "reason": "表头和代表性业务行共同支持该字段映射",
        }

    model = AIModelInjectionWorkbookMappingV1.model_validate(
        {
            "document_kind": "PLANNED_SCHEDULE",
            "source_sha256": source_hash,
            "source_sheet": {
                "sheet_name": "1月",
                "header_rows": [2, 3],
                "data_start_row": 4,
                "data_end_row": 8,
            },
            "row_layout": {
                "layout_type": "GROUPED_BY_MACHINE",
                "machine_code_strategy": "CURRENT_OR_INHERITED",
                "machine_header_rule": "SAME_VALUE_IN_TWO_COLUMNS",
                "machine_header_columns": ["A", "B"],
                "task_identity_fields": ["mold_no", "order_no"],
                "backlog_rule": "BUSINESS_ROW_WITHOUT_MACHINE",
                "termination": {
                    "mode": "FIRST_FOOTER_LABEL",
                    "footer_labels": ["合计"],
                },
            },
            "field_mappings": [
                field("machine_code", "B", "identifier"),
                field("mold_no", "G", "identifier", header_cell),
                field("order_no", "I", "identifier"),
                field("order_quantity", "L", "number"),
                field("completed_quantity", "M", "number"),
                field("planned_start", "AG", "datetime"),
                field("planned_finish", "AH", "datetime"),
            ],
            "shift_grid": {"enabled": False},
            "warnings": [],
            "overall_confidence": 0.95,
        }
    )
    payload = model.model_dump(mode="json")
    digest = hashlib.sha256(
        json.dumps(
            payload, ensure_ascii=False, separators=(",", ":"), sort_keys=True
        ).encode()
    ).hexdigest()
    return AIInjectionWorkbookMappingV1(
        **payload, generated_by_model="qwen3.7-plus", mapping_digest=digest
    )


class _SequenceProvider:
    def __init__(self, responses: list[ProviderResponse]) -> None:
        self.responses = responses
        self.requests = []

    async def generate(self, request):
        self.requests.append(request)
        return self.responses.pop(0)

    async def aclose(self) -> None:
        return None


def _provider_response(payload: dict[str, object], call_id: str) -> ProviderResponse:
    return ProviderResponse(
        tool_calls=(
            ProviderToolCall(
                call_id=call_id,
                name="workbook.submit_injection_workbook_mapping",
                arguments_json=json.dumps(payload, ensure_ascii=False),
            ),
        ),
        response_id=f"response-{call_id}",
    )


def test_unified_schema_accepts_demand_without_machine_or_shift() -> None:
    mapping = AIModelInjectionWorkbookMappingV1.model_validate(
        {
            "document_kind": "DEMAND_ORDER",
            "source_sha256": "a" * 64,
            "source_sheet": {
                "sheet_name": "下单表",
                "header_rows": [3],
                "data_start_row": 4,
                "data_end_row": 10,
            },
            "row_layout": {
                "layout_type": "FLAT_ROWS",
                "machine_code_strategy": "NONE",
                "machine_header_rule": "NONE",
                "task_identity_fields": ["source_mold_no", "product_name"],
            },
            "field_mappings": [
                {
                    "canonical_field": field,
                    "source_column": column,
                    "header_cell": f"{column}3",
                    "transformer": transformer,
                    "confidence": 0.9,
                    "reason": "表头与需求业务行共同支持映射",
                }
                for field, column, transformer in (
                    ("source_mold_no", "A", "identifier"),
                    ("product_name", "B", "trim"),
                    ("order_quantity", "C", "number"),
                )
            ],
            "shift_grid": {"enabled": False},
            "overall_confidence": 0.9,
        }
    )
    assert mapping.document_kind == "DEMAND_ORDER"


def test_structural_signature_ignores_month_values_order_values_and_day_count() -> None:
    first = _workbook(month=1, shift_days=2)
    second = _workbook(month=2, shift_days=3)
    first_packet = build_workbook_recognition_packet(
        source_file_name="first.xlsx",
        content=first,
        factory_id="huaxing",
        business_date="2026-01-01",
    )
    second_packet = build_workbook_recognition_packet(
        source_file_name="second.xlsx",
        content=second,
        factory_id="huaxing",
        business_date="2026-02-01",
    )
    assert structural_layout_signature(first_packet) == structural_layout_signature(
        second_packet
    )
    sampled_rows = {row.row for row in first_packet.sheets[0].representative_rows}
    assert {4, 5, 8} <= sampled_rows


def test_mapping_source_validation_rejects_fabricated_header_cell() -> None:
    content = _workbook()
    packet = build_workbook_recognition_packet(
        source_file_name="plan.xlsx",
        content=content,
        factory_id="huaxing",
        business_date="2026-01-01",
    )
    validate_mapping_sources(
        _mapping(packet.source.source_sha256),
        packet,
        requested_kind="PLANNED_SCHEDULE",
    )
    with pytest.raises(ValueError, match="表头单元格"):
        validate_mapping_sources(
            _mapping(packet.source.source_sha256, header_cell="G9"),
            packet,
            requested_kind="PLANNED_SCHEDULE",
        )


def test_unified_skill_uses_one_controlled_repair_and_required_tool() -> None:
    content = _workbook()
    packet = build_workbook_recognition_packet(
        source_file_name="plan.xlsx",
        content=content,
        factory_id="huaxing",
        business_date="2026-01-01",
    )
    valid = _mapping(packet.source.source_sha256).model_dump(
        mode="json",
        exclude={
            "generated_by_model",
            "skill_id",
            "skill_version",
            "prompt_version",
            "mapping_digest",
        },
    )
    invalid = json.loads(json.dumps(valid))
    invalid["source_sheet"]["sheet_name"] = "不存在"
    provider = _SequenceProvider(
        [
            _provider_response(invalid, "invalid"),
            _provider_response(valid, "valid"),
        ]
    )
    settings = Settings(
        _env_file=None,
        ai_enabled=True,
        ai_provider="fake",
        ai_default_model="qwen3.7-plus",
        qwen_table_model="qwen3.7-plus",
        ai_cloud_workbook_mapping_enabled=True,
    )
    result = asyncio.run(
        recognize_injection_workbook_mapping(
            packet=packet,
            provider=provider,
            settings=settings,
            request_id="unified-mapping-test",
            requested_kind="PLANNED_SCHEDULE",
        )
    )
    assert result.skill_id == "injection_scheduling.workbook_mapping"
    assert len(provider.requests) == 2
    assert provider.requests[0].tools[0].name == (
        "workbook.submit_injection_workbook_mapping"
    )
    assert provider.requests[0].tool_choice_policy.value == "REQUIRED"
    assert provider.requests[0].data_classification.value == "CONFIDENTIAL"
    assert "<UNTRUSTED_WORKBOOK_PACKET>" in provider.requests[0].input[1].content
    assert "唯一一次修复机会" in provider.requests[1].input[1].content
