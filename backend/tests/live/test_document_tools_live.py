from __future__ import annotations

import os
import sys
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

BACKEND_DIR = Path(__file__).resolve().parents[2]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.core.config import settings
from app.services.document_studio.renderers.office_pdf_renderer import (
    render_docx_to_pdf,
)
from app.services.pdf_split import split_pdf
from app.services.pdf_to_excel import convert_pdf_to_excel
from app.services.pdf_to_word import convert_pdf_to_word

pytestmark = pytest.mark.skipif(
    os.getenv("RUN_LIVE_DOCUMENT_TESTS") != "1",
    reason="set RUN_LIVE_DOCUMENT_TESTS=1 to run real document dependencies",
)


def _fixture(name: str) -> tuple[bytes, str]:
    raw = os.getenv(name, "").strip()
    if not raw:
        pytest.fail(f"{name} must point to a real test document")
    path = Path(raw)
    if not path.is_file():
        pytest.fail(f"{name} does not exist: {path}")
    return path.read_bytes(), path.name


def test_real_text_pdf_to_excel_reopens() -> None:
    data, filename = _fixture("DOCUMENT_LIVE_TEXT_PDF")
    result = convert_pdf_to_excel(data, filename)
    workbook = load_workbook(BytesIO(result.content))
    assert len(workbook.worksheets) >= 1


def test_real_scanned_pdf_to_word_reopens() -> None:
    data, filename = _fixture("DOCUMENT_LIVE_SCAN_TABLE_PDF")
    result = convert_pdf_to_word(data, filename)
    Document(BytesIO(result.content))
    assert result.page_count >= 1


def test_real_docx_uses_installed_libreoffice_and_reopens() -> None:
    data, filename = _fixture("DOCUMENT_LIVE_DOCX")
    Document(BytesIO(data))
    result = render_docx_to_pdf(
        data,
        filename,
        command=settings.document_office_renderer_command,
        timeout_seconds=settings.document_office_renderer_timeout_seconds,
        network_isolation_command=None,
    )
    assert len(PdfReader(BytesIO(result.content)).pages) >= 1


def test_real_pdf_ranges_split_to_safe_reopenable_zip() -> None:
    data, filename = _fixture("DOCUMENT_LIVE_SPLIT_PDF")
    source_pages = len(PdfReader(BytesIO(data)).pages)
    assert source_pages >= 2
    result = split_pdf(data, filename, mode="ranges", page_ranges="1, 2")
    with ZipFile(BytesIO(result.content)) as archive:
        names = archive.namelist()
        assert len(names) == 2
        assert all(
            ".." not in name and not name.startswith(("/", "\\")) for name in names
        )
        assert all(
            len(PdfReader(BytesIO(archive.read(name))).pages) == 1 for name in names
        )
