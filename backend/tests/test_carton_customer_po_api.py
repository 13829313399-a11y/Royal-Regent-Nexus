from types import SimpleNamespace

import pytest
from test_molding_sample_api import make_client, login_as
from test_carton_master_api import prepare, create
from test_carton_procurement_api import _order_payload
from test_carton_history_preview_api import upload, BASE
from test_carton_history_identity_api import _customer, _orders
from test_carton_history_wide import workbook, row, MINIMAL_HEADERS
from app.services.carton_master import number_warnings
from app.services.carton_procurement_imports import _match_rows


def test_customer_po_orders_keep_normal_submission_flow(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        first = create(client, customer_po="PO-100")
        second = create(client, customer_po="PO-101")
        assert first["status"] == second["status"] == "CONFIRMED"
        assert client.post(BASE + "/orders", json={**_order_payload(), "customer_po":"PO-100"}).status_code == 409
        url = BASE + f"/orders/{first['order_no']}/offline-supplement"
        assert client.post(url, json={"factory_id":"huaxing", "expected_revision":first["revision"], "posting_confirmed":True}).status_code == 404
        legacy = client.post(BASE + "/orders", json={**_order_payload(), "customer_po":"PO-102", "offline_placed":True, "posting_confirmed":True})
        # Unknown old flags may be rejected or ignored, but can never bypass normal confirmation.
        assert legacy.status_code == 422 or (legacy.status_code == 201 and legacy.json()["status"] == "CONFIRMED")
        assert client.get(BASE + "/inventory/movements", params={"factory_id":"huaxing"}).json()["total"] == 0


def test_history_groups_and_deduplicates_by_customer_po(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin"); _customer(client, "迪奇")
        content = workbook([row(客户PO="P001"), row(客户PO="P002")], MINIMAL_HEADERS + ["客户PO"])
        preview = upload(client, content).json()
        assert not preview["errors"], preview
        assert preview["group_count"] == 2
        result = upload(client, content, preview=False, fingerprint=preview["source_fingerprint"])
        assert result.status_code == 201, result.text
        assert {o["customer_po"] for o in _orders(client)} == {"P001", "P002"}
        assert upload(client, content).json()["skipped_count"] == 2
        explicit = workbook([row(客户PO="P001", 历史订单号="NEW-HISTORY-ID")], MINIMAL_HEADERS + ["客户PO", "历史订单号"])
        assert upload(client, explicit).json()["skipped_count"] == 1


def test_blank_po_remains_optional_even_with_block_rule():
    rules = {"customer_po_rule": {"mode":"BLOCK", "prefix":"PO-", "min_length":6, "max_length":20}}
    assert not number_warnings(rules, "C", "I", "")
    assert number_warnings(rules, "C", "I", "WRONG")[0]["blocking"]
    assert not number_warnings(rules, "C", "I", "PO-123")


@pytest.mark.parametrize("po,expected", [("", "AMBIGUOUS"), ("MISSING", "MISSING_ORDER"), ("P002", "MATCHED")])
def test_import_matching_never_uses_other_po(po, expected):
    # Material narrowing must not guess which customer's purchase order a row belongs to.
    pairs=[]
    for index in (1,2):
        order=SimpleNamespace(id=f"O{index}", order_no=f"O{index}", contract_no="C", item_no="I", customer_po=f"P00{index}", customer_code="D", customer_name="D", status="PENDING_SUPPLIER", due_date="2026-09-20", product_order_quantity=100)
        line=SimpleNamespace(id=f"L{index}", order_id=order.id, packaging_type="外箱", paper_quality="A", specification="1*2*3", required_quantity=10, unit="个", unit_price=2)
        pairs.append((line,order))
    class FakeDB:
        def execute(self, _):
            return SimpleNamespace(all=lambda:pairs)
    rows=[{"contract_no":"C", "item_no":"I", "customer_po":po, "delivered_quantity":1}]
    _match_rows(FakeDB(), "huaxing", "DELIVERY_NOTE", rows)
    assert rows[0]["match_status"] == expected
    if po == "P002":
        assert rows[0]["order_line_id"] == "L2"


@pytest.mark.parametrize("actual, supplied, expected", [("PO-1", "PO1", "MISSING_ORDER"), ("客户甲-1", "客户乙-1", "MISSING_ORDER"), ("客户甲-1", "客户甲-1", "MATCHED")])
def test_po_matching_preserves_punctuation_and_chinese(actual, supplied, expected):
    order = SimpleNamespace(id="O", order_no="O", contract_no="C", item_no="I", customer_po=actual,
        customer_code="D", customer_name="D", status="PENDING_SUPPLIER", due_date="2026-09-20")
    line = SimpleNamespace(id="L", order_id="O", packaging_type="外箱", paper_quality="A", specification="1*2*3", required_quantity=10, unit="个", unit_price=2)
    class FakeDB:
        def execute(self, _): return SimpleNamespace(all=lambda:[(line,order)])
    rows=[{"contract_no":"C", "item_no":"I", "customer_po":supplied, "delivered_quantity":1}]
    _match_rows(FakeDB(), "huaxing", "DELIVERY_NOTE", rows)
    assert rows[0]["match_status"] == expected


def test_customer_po_rule_is_saved_and_enforced_by_backend(monkeypatch):
    with make_client(monkeypatch) as client:
        prepare(client)
        result=client.post(BASE+"/master-data",json={"factory_id":"huaxing","kind":"RULE","customer_code":"DICKIE",
            "data":{"customer_po_rule":{"mode":"BLOCK","prefix":"PO-","min_length":6}},"reason":"客户要求核对采购单号"})
        assert result.status_code == 201, result.text
        assert client.post(BASE+"/orders",json={**_order_payload(),"customer_po":"BAD"}).status_code == 422
        assert create(client)["customer_po"] == ""
        assert create(client, customer_po="PO-001")["customer_po"] == "PO-001"


def test_receiving_one_po_does_not_receive_the_other(monkeypatch):
    from test_carton_direct_receipt_api import payload
    from test_carton_procurement_api import _submit_order
    with make_client(monkeypatch) as client:
        prepare(client)
        first = _submit_order(client, create(client, customer_po="PO-A"))
        second = _submit_order(client, create(client, customer_po="PO-B"))
        result = client.post(BASE + "/receipts", json=payload(first))
        assert result.status_code == 201, result.text
        orders = {o["customer_po"]: o for o in _orders(client)}
        assert orders["PO-A"]["status"] == "COMPLETED"
        assert orders["PO-B"]["status"] == "PENDING_SUPPLIER"
        assert all(float(line["received_quantity"]) == 0 for line in orders["PO-B"]["lines"])
        movements = client.get(BASE + "/inventory/movements", params={"factory_id":"huaxing"}).json()["items"]
        assert {m["order_line_id"] for m in movements} == {line["id"] for line in first["lines"]}
        assert not {line["id"] for line in second["lines"]} & {m["order_line_id"] for m in movements}


def test_vertical_legacy_history_preserves_customer_po():
    from io import BytesIO
    from openpyxl import load_workbook
    from app.services.carton_procurement_history_import import _parse_rows
    from test_carton_procurement_api import _history_workbook_bytes
    from test_carton_history_identity_api import _row
    wb = load_workbook(BytesIO(_history_workbook_bytes([_row("迪奇")])))
    ws = wb.active
    column = ws.max_column + 1
    ws.cell(5, column, "客户PO")
    ws.cell(6, column, "客户甲-001")
    stream = BytesIO(); wb.save(stream)
    rows, _ = _parse_rows("legacy.xlsx", stream.getvalue())
    assert rows[0]["customer_po"] == "客户甲-001"
