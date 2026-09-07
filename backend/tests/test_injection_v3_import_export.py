import io
import os
from pathlib import Path
from uuid import uuid4

import pytest
from app.models.injection_scheduling import Demand, HistoricalOutput, Machine
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session
from test_injection_v3_api import BASE, ok, seed_job, write
from test_injection_v3_api import env as _env_fixture

env = _env_fixture


def upload(env, content, sheet="计划表", factory="huaxing"):
    revision = (
        env[0].get(BASE + "/summary", params={"factory_id": factory}).json()["revision"]
    )
    return env[0].post(
        BASE + "/imports",
        data={
            "factory_id": factory,
            "base_revision": revision,
            "client_operation_id": uuid4().hex,
            "sheet_name": sheet,
        },
        files={
            "file": (
                "plan.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


def test_export_roundtrip_ids_window_and_edit(env):
    _, _, demand = seed_job(env)
    downloaded = env[0].post(
        BASE + "/exports/plan",
        json={"factory_id": "huaxing", "window_start": "2026-09-01", "window_days": 14},
    )
    assert downloaded.status_code == 200, downloaded.text
    book = load_workbook(io.BytesIO(downloaded.content), data_only=False)
    assert book.sheetnames == ["计划表", "标准交换表", "口径说明"]
    assert book["计划表"]["A1"].value.startswith("华兴")
    exchange = book["标准交换表"]
    headers = {exchange.cell(3, c).value: c for c in range(1, exchange.max_column + 1)}
    assert exchange.cell(4, headers["id"]).value == demand["id"]
    assert exchange.cell(4, headers["window_opening_shots"]).value == 0
    preview = ok(upload(env, downloaded.content, "标准交换表"))
    applied = ok(write(env, f"/imports/{preview['batch_id']}/apply"))
    assert applied["summary"]["unchanged_count"] == 1, applied["summary"]
    assert (
        env[0]
        .get(BASE + "/summary", params={"factory_id": "huaxing"})
        .json()["total_count"]
        == 1
    )
    # Editing a writable field in the same stable-ID exchange changes one demand.
    exchange.cell(4, headers["planned_shots"], 1200)
    output = io.BytesIO()
    book.save(output)
    preview = ok(upload(env, output.getvalue(), "标准交换表"))
    applied = ok(write(env, f"/imports/{preview['batch_id']}/apply"))
    assert applied["summary"]["applied_count"] == 1
    assert (
        env[0]
        .get(BASE + f"/demands/{demand['id']}", params={"factory_id": "huaxing"})
        .json()["demand"]["planned_shots"]
        == 1200
    )
    assert upload(env, downloaded.content, "标准交换表", "huadeng").status_code == 422


def test_real_import_persistence_and_history_not_double_counted(env):
    path = os.getenv("INJECTION_V3_WORKBOOK")
    if not path:
        pytest.skip("Set INJECTION_V3_WORKBOOK for private workbook integration")
    content = Path(path).read_bytes()
    preview = ok(upload(env, content))
    assert preview["summary"]["machine_count"] == 76
    applied = ok(write(env, f"/imports/{preview['batch_id']}/apply"))
    assert applied["summary"]["applied_count"] == 280
    with Session(env[1]) as db:
        assert db.scalar(select(func.count()).select_from(Machine)) == 76
        assert db.scalar(select(func.count()).select_from(Demand)) == 280
        assert db.scalar(select(func.count()).select_from(HistoricalOutput)) == 1449
        assert db.scalar(select(func.sum(Demand.completed_shots))) == 2253660
        assert db.scalar(select(func.sum(HistoricalOutput.quantity))) == 2267721
        tail = db.scalar(select(Demand).where(Demand.source_row == 360))
        assert tail.remaining_shots == 100000 and tail.machine_code is None
    repeated = ok(upload(env, content))
    assert repeated["batch_id"] == preview["batch_id"]
    second = ok(write(env, f"/imports/{preview['batch_id']}/apply"))
    assert second["already_applied"] is True
