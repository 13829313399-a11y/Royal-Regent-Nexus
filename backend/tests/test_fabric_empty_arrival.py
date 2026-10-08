"""Pending workbook emptiness is a business fact; arbitrary errors remain unknown."""
import importlib
import json
from io import BytesIO
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from test_fabric_procurement import BASE, book, row, upload, commit, stored
from test_fabric_procurement_tracking import save
from test_fabric_receiving import request, receive
from test_molding_sample_api import make_client, login_as


def formula_book(*, cached=None, reverse=False, formula=None, detail="", returned=False):
    data = row(入库数量=None, 交货明细=detail)
    wb = load_workbook(BytesIO(book(None if returned else [data], [data] if returned else None, reverse=reverse)))
    ws = wb["已回料" if returned else "未回物料"]
    header_row = 1 if returned else 2
    headers = {cell.value: cell.column for cell in ws[header_row]}
    n = header_row + 1
    address = f"{get_column_letter(headers['入库数量'])}{n}"
    detail_address = f"{get_column_letter(headers['交货明细'])}{n}"
    ws[address] = formula if formula is not None else f"=EVALUATE({detail_address})"
    out = BytesIO(); wb.save(out)
    if cached is None:
        return out.getvalue()
    result = BytesIO()
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(out) as src, ZipFile(result, "w") as dest:
        for member in src.infolist():
            value = src.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(value)
                cell = root.find(f".//s:c[@r='{address}']", ns)
                cell.set("t", "e")
                cell.find("s:v", ns).text = cached
                value = ET.tostring(root, encoding="utf-8")
            dest.writestr(member, value)
    return result.getvalue()


@pytest.mark.parametrize("reverse,cached", [(False, None), (True, None), (False, "#VALUE!"), (True, "#VALUE!")])
def test_empty_detail_evaluate_cache_is_zero_only_for_its_own_pending_row(reverse, cached):
    parser = importlib.import_module("app.services.fabric_procurement_parser")
    parsed = parser.parse_workbook("source.xlsx", formula_book(reverse=reverse, cached=cached))
    entry = parsed["rows"][0]
    assert entry["facts"]["reported_received_quantity"] == "0"
    assert entry["facts"]["reported_received_quantity_basis"] == "EMPTY_PENDING_DELIVERY"
    assert entry["raw"]["formulas"] and entry["raw"]["headers"]
    if cached:
        received_col = next(k for k, v in entry["raw"]["headers"].items() if v == "reported_received_quantity")
        assert entry["raw"]["values"][received_col] == "#VALUE!"
    assert not entry["errors"]


@pytest.mark.parametrize("content", [
    book([row(入库数量="#VALUE!", 交货明细="")]),
    formula_book(formula="=EVALUATE(L4)", cached="#VALUE!"),
    formula_book(formula="=SUM(L3)", cached="#VALUE!"),
    formula_book(detail="=SUM(A3:A4)"),
    book([row(入库数量=None, 交货明细="一卷已到，数量待称")]),
    book([row(入库数量=50, 交货明细="20+20")]),
    book(None, [row(入库数量=None, 交货明细="")]),
    formula_book(returned=True, cached="#VALUE!"),
], ids=["literal-error", "other-row", "other-formula", "detail-formula", "unclear-detail", "conflict", "returned-blank", "returned-formula"])
def test_other_unknowns_and_returned_sheets_are_not_zeroed(content):
    parser = importlib.import_module("app.services.fabric_procurement_parser")
    facts = parser.parse_workbook("source.xlsx", content, scope="ALL")["rows"][0]["facts"]
    assert facts["reported_received_quantity"] is None
    assert "reported_received_quantity_basis" not in facts


def test_missing_quantity_columns_and_future_adapter_do_not_infer_empty_arrival():
    parser = importlib.import_module("app.services.fabric_procurement_parser")
    values = {"order_no": "P1", "supplier": "S", "material_code": "M", "material_name": "布", "unit": "码",
              "ordered_quantity": 100, "reported_received_quantity": None, "delivery_detail": ""}
    assert parser.normalize_row(values, "PENDING")[0]["reported_received_quantity"] is None
    wb = load_workbook(BytesIO(book([row(入库数量=None, 交货明细="")])))
    ws = wb["未回物料"]
    ws.cell(2, 12, "其他资料")
    out = BytesIO(); wb.save(out)
    assert parser.parse_workbook("source.xlsx", out.getvalue())["rows"][0]["facts"]["reported_received_quantity"] is None


def shared_book(base_formula="EVALUATE(L3)", *, index="1", duplicate=False):
    wb = load_workbook(BytesIO(book([row(入库数量=None, 交货明细=""), row(生产单号="P002", 入库数量=None, 交货明细="")])))
    wb["未回物料"]["M3"] = "=" + base_formula
    wb["未回物料"]["M4"] = "=EVALUATE(L4)"
    out = BytesIO(); wb.save(out); result = BytesIO()
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(out) as src, ZipFile(result, "w") as dest:
        for member in src.infolist():
            value = src.read(member.filename)
            if member.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(value)
                for address in ("M3", "M4"):
                    cell = root.find(f".//s:c[@r='{address}']", ns)
                    formula = cell.find("s:f", ns)
                    formula.set("t", "shared")
                    if index is not None: formula.set("si", index)
                    if address == "M3": formula.set("ref", "M3:M4")
                    elif duplicate: formula.text = "EVALUATE(L4)"; formula.set("ref", "M3:M4")
                    else: formula.text = None
                    cell.set("t", "e"); cell.find("s:v", ns).text = "#VALUE!"
                value = ET.tostring(root, encoding="utf-8")
            dest.writestr(member, value)
    return result.getvalue()


@pytest.mark.parametrize("base_formula,expected", [("EVALUATE(L3)", "0"), ("EVALUATE($L3)", "0"),
    ("EVALUATE(L$3)", None), ("SUM(L3)", None), ("EVALUATE(K3)", None)])
def test_shared_formula_tail_resolves_only_a_bounded_same_column_relative_reference(base_formula, expected):
    parser = importlib.import_module("app.services.fabric_procurement_parser")
    tail = parser.parse_workbook("source.xlsx", shared_book(base_formula))["rows"][1]
    assert tail["facts"]["reported_received_quantity"] == expected
    assert tail["raw"]["formulas"]["M4"] == "" and tail["raw"]["formula_metadata"]["M4"]["si"] == "1"


@pytest.mark.parametrize("index", [None, "bad", "-1", "4294967296", "１", "١"])
def test_invalid_shared_index_does_not_establish_zero(index):
    parser = importlib.import_module("app.services.fabric_procurement_parser")
    assert parser.parse_workbook("source.xlsx", shared_book(index=index))["rows"][1]["facts"]["reported_received_quantity"] is None


def test_duplicate_shared_anchors_block_the_workbook_before_persisting_any_rows():
    parser = importlib.import_module("app.services.fabric_procurement_parser")
    with pytest.raises(HTTPException) as exc:
        parser.parse_workbook("source.xlsx", shared_book(duplicate=True))
    assert exc.value.status_code == 422 and "共享公式分组重复" in exc.value.detail


def test_full_order_chase_decrements_actual_receipts_once_even_after_procurement_updates(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        initial = book([row(采购明细ID="BLANK", 入库数量=None, 交货明细="")])
        save(client, initial)
        line = stored(client)["items"][0]
        assert line["starting_chase_quantity"] == line["warehouse_outstanding_quantity"] == "100"
        assert line["prior_received_quantity"] == "0" and not line["receipt_quantity_review_required"]
        posted = receive(client, line, request(line, batches=[{"quantity": "10", "location": "A1", "dye_lot": "0001"}]))
        assert posted.status_code == 200, posted.text
        assert posted.json()["prior_received_quantity"] == "0"
        save(client, book([row(采购明细ID="BLANK", 入库数量=10, 交货明细="10")]))
        line = stored(client)["items"][0]
        assert line["starting_chase_quantity"] == "100" and line["warehouse_outstanding_quantity"] == "90"
        save(client, initial)
        line = stored(client)["items"][0]
        assert line["warehouse_outstanding_quantity"] == "90" and line["warehouse_received_quantity"] == "10"
        assert receive(client, line, request(line, delivery_reference="DN2", batches=[{"quantity": "90", "location": "A1", "dye_lot": "0002"}])).status_code == 200
        assert stored(client, view="OUTSTANDING")["total"] == 0
        assert client.get(BASE + "/receipts", params={"factory_id": "huakang-c"}).json()["total"] == 2
        assert client.get(BASE + "/stock", params={"factory_id": "huakang-c"}).json()["total"] == 2


def test_same_purchase_different_materials_are_separate_detail_quantities(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        save(client, book([row(订单号="CGDD014333", 订单数量=65, 入库数量=53, 交货明细="53")],
                          [row(订单号="CGDD014333", 物料编码="OTHER", 订单数量=19, 入库数量=19, 交货明细="19")]), scope="ALL")
        items = stored(client)["items"]
        assert len(items) == 2
        assert {line["facts"]["material_code"]: line["starting_chase_quantity"] for line in items} == {"000123": "12", "OTHER": "0"}
        assert stored(client, view="OUTSTANDING")["total"] == 1
        assert client.get(BASE + "/stock", params={"factory_id": "huakang-c"}).json()["total"] == 0


def test_legacy_duplicate_rows_upgrade_only_with_exact_evidence_and_keep_ids(monkeypatch):
    with make_client(monkeypatch) as client:
        parser = importlib.import_module("app.services.fabric_procurement_parser")
        detect = parser.pending_empty_reference
        login_as(client, "admin")
        content = book([row(入库数量=None, 交货明细="", 单价=1), row(入库数量=None, 交货明细="", 单价=2)])
        monkeypatch.setattr(parser, "pending_empty_reference", lambda *args: False)
        save(client, content)
        old = {line["facts"]["unit_price"]: line["id"] for line in stored(client)["items"]}
        monkeypatch.setattr(parser, "pending_empty_reference", detect)
        changed = book([row(入库数量=None, 交货明细="", 单价=3), row(入库数量=None, 交货明细="", 单价=2)])
        assert upload(client, changed).json()["counts"]["blocked"] == 2
        preview = upload(client, content).json()
        assert preview["counts"]["updated"] == 2 and preview["counts"]["blocked"] == 0
        assert commit(client, content, preview).status_code == 200
        lines = stored(client)["items"]
        assert {line["facts"]["unit_price"]: line["id"] for line in lines} == old
        assert all(line["starting_chase_quantity"] == "100" for line in lines)
        assert upload(client, content).json()["counts"]["unchanged"] == 2
        dbm = importlib.import_module("app.db")
        model = importlib.import_module("app.models.fabric_procurement")
        with dbm.SessionLocal() as db:
            evidence = list(db.scalars(select(model.FabricProcurementEvidence)))
            assert len(evidence) == 4
            assert sum(json.loads(event.after_json)["reported_received_quantity"] is None for event in evidence) == 2


def test_mixed_legacy_duplicates_reserve_proven_upgrade_before_unchanged_matching(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        parser = importlib.import_module("app.services.fabric_procurement_parser")
        detect = parser.pending_empty_reference
        wb = load_workbook(BytesIO(book([row(入库数量=None, 交货明细=""), row(入库数量=None, 交货明细="")])))
        ws = wb["未回物料"]
        ws["M3"] = "=EVALUATE(L3)"; ws["M4"] = "=SUM(L4)"
        out = BytesIO(); wb.save(out); content = out.getvalue()
        monkeypatch.setattr(parser, "pending_empty_reference", lambda *args: False)
        save(client, content)
        dbm = importlib.import_module("app.db")
        model = importlib.import_module("app.models.fabric_procurement")
        with dbm.SessionLocal() as db:
            ids = {event.row_number: event.line_id for event in db.scalars(select(model.FabricProcurementEvidence))}
        monkeypatch.setattr(parser, "pending_empty_reference", detect)
        preview = upload(client, content).json()
        assert preview["counts"]["updated"] == 1 and preview["counts"]["unchanged"] == 1 and preview["counts"]["blocked"] == 0
        assert commit(client, content, preview).status_code == 200
        lines = {line["id"]: line for line in stored(client)["items"]}
        assert lines[ids[3]]["starting_chase_quantity"] == "100"
        assert lines[ids[4]]["starting_chase_quantity"] is None
        with dbm.SessionLocal() as db:
            assert len(list(db.scalars(select(model.FabricProcurementEvidence)))) == 3
