from io import BytesIO

import openpyxl
import pytest

from app.services.huadeng_order_legacy.goliath_po_parser import parse_po
from app.services import customer_order_unified as service


def po(*,price_tail='3',total='66436.80',brand='Goliath Far East Ltd',with_total=True):
    commands=[]
    def text(x,y,v):
        escaped=v.replace('\\','\\\\').replace('(','\\(').replace(')','\\)')
        commands.append(f'BT /F1 8 Tf {x} {842-y} Td ({escaped}) Tj ET')
    for x,y,v in [(50,50,'Purchase Order'),(300,70,brand),(300,175,'Deliver to :'),(300,190,'Pressman Toy Corp. c/o Regal Logistics'),(50,260,'Order # T_10130 - PHK26_00694'),(50,275,'Order date 24 April 2026'),(50,348,'Item'),(106,348,'Description'),(282,348,'ITF14'),(306,348,'Barcode'),(385,348,'No.'),(432,348,'Qty'),(463,348,'price'),(509,348,'Amount'),(55,368,'935373.002'),(112,368,'Butt Face - Chester Cheeks'),(271,368,'18720077353739'),(355,368,'935373.002'),(409,368,'16,000'),(456,368,'4.152'),(472,379,price_tail),(497,368,total),(55,391,'USDOMPO0'),(112,415,'CARGO READY DATE:'),(112,430,'15-AUG-2026'),(112,460,'US PO NO.: PUS26_00435'),(320,575,'Total USD Excl. VAT '+total)]:text(x,y,v)
    if not with_total:commands.pop()
    return _pdf_document(commands)


def _pdf_document(commands):
    stream='\n'.join(commands).encode('ascii')
    objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [4 0 R] /Count 1 >>',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R >> >> /Contents 5 0 R >>',f'<< /Length {len(stream)} >>\nstream\n'.encode()+stream+b'\nendstream']
    content=b'%PDF-1.4\n';offsets=[]
    for i,obj in enumerate(objects,1):
        offsets.append(len(content));content+=f'{i} 0 obj\n'.encode()+obj+b'\nendobj\n'
    start=len(content);content+=b'xref\n0 6\n0000000000 65535 f \n'
    content+=b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets)
    return content+f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF'.encode()


def template():
    w=openpyxl.Workbook();w.remove(w.active)
    for name,marker in ((service.ITEM_SHEET,7),(service.ORDER_SHEET,6),(service.REVIEW_SHEET,8)):
        s=w.create_sheet(name)
        for c,v in enumerate(service.ITEM_HEADERS if name==service.ITEM_SHEET else service.ORDER_HEADERS,1):s.cell(3,c,v)
        s.cell(marker,1,'取消单')
    s=w[service.ITEM_SHEET]
    for c,v in {4:'HIST',5:'HIST',7:'Walmart',8:'935373.002',9:'中文品名',11:500,12:2}.items():s.cell(4,c,v)
    b=BytesIO();w.save(b);return b.getvalue()


def test_pdf_columns_keep_wrapped_price():
    rows=parse_po(po(),'real.pdf')
    assert len(rows)==1
    r=rows[0]
    assert r['unit_price_usd']=='4.1523'
    assert r['amount_usd']=='66436.80'
    assert r['customer']=='Pressman Toy Corp. c/o Regal Logistics'
    assert r['item_no']=='935373.002'
    assert r['flags']==[]


def test_inconsistent_price_is_explicit_blocker():
    r=parse_po(po(price_tail='2'),'bad.pdf')[0]
    assert r['flags'][0]['code']=='price_amount_conflict'


def test_wrong_customer_rejected():
    with pytest.raises(ValueError,match='不是可识别'):
        parse_po(po(brand='Unrelated supplier'),'wrong.pdf')


def test_missing_total_cannot_silently_pass_partial_document():
    with pytest.raises(ValueError,match='唯一可核对'):
        parse_po(po(with_total=False),'partial.pdf')


def test_latest_workbook_export_identity_history_and_duplicate_policy():
    args=dict(factory_id='huadeng',customer_code='goliath',received_date='2026-09-07',po_files=[('po.pdf',po())],schedule_file_name='latest.xlsx',schedule_content=template())
    output,_,preview=service.export_unified_customer_schedule(**args)
    r=preview['rows'][0]
    assert r['reference_no']==r['contract_no']=='T_10130-PHK26_00694'
    assert r['po_no']=='PUS26_00435'
    assert r['customer_name'].startswith('Pressman')
    assert r['product_name_zh']=='中文品名'
    assert r['units_per_carton']=='2'
    assert r['unit_price_hkd']=='32.180325'
    assert r['date_code']=='RR2176'
    w=openpyxl.load_workbook(BytesIO(output))
    assert w[service.ITEM_SHEET]['L5'].value==2
    assert w[service.ITEM_SHEET]['Y5'].value is None
    assert w[service.ORDER_SHEET]['J4'].value=='=IF(LEN(\'ITEM表\'!K5)=0,"",\'ITEM表\'!K5)'
    assert w[service.ORDER_SHEET]['P4'].value==4.1523
    assert w[service.ORDER_SHEET]['N4'].value=='=P4*7.75'
    repeated=service.create_unified_customer_preview(**{**args,'schedule_content':output})
    assert any(i['code']=='duplicate_existing_order' for i in repeated['rows'][0]['issues'])
    with pytest.raises(service.CustomerOrderUnifiedError):
        service.create_unified_customer_preview(**{**args,'factory_id':'huakang-d'})


def european_po(*, note_date='20-JAN-2025', pack='24', unknown_layout=False, date_label='SHIP DATE:'):
    commands=[]
    entries=[
        (50,50,'Purchase Order'), (300,70,'Goliath BV'),
        (300,175,'Deliver to :'), (300,190,'Schotpoort'),
        (50,260,'Order # T_10130 - PEU25_00642'), (50,275,'Order date 26 September 2025'),
        (50,348,'Item'), (106,348,'Description'),
        (269,337,'Requested'), (269,348,'delivery date'),
        (381,337,'Qty in a'), (376,348,'Unknown' if unknown_layout else 'Casepack'),
        (432,348,'Qty'), (463,348,'price'), (509,348,'Amount'),
        (50,368,'934811.024'), (106,368,"FLIP 'N TRiX 2x12-CDU (12L)"),
        (269,368,'20 January 2026'), (401,368,pack), (421,368,'30,000'),
        (461,368,'1.094'), (502,368,'32,820.00'),
        (50,390,'EUDOMPO001'), (269,390,'20 January 2026'),
        (106,430,date_label), (106,445,note_date),
        (50,480,'POAMEND'), (269,480,'20 January 2026'),
        (106,495,'THIS IS 3RD REVISION.'),
        (320,575,'Total USD Excl. VAT 32,820.00'),
    ]
    for x,y,value in entries:
        escaped=value.replace('\\','\\\\').replace('(','\\(').replace(')','\\)')
        commands.append(f'BT /F1 8 Tf {x} {842-y} Td ({escaped}) Tj ET')
    return _pdf_document(commands)


def test_european_layout_uses_explicit_casepack_and_product_delivery_date():
    row=parse_po(european_po(),'PEU25_00642R3.pdf')[0]
    assert row['contract_no']=='T_10130-PEU25_00642'
    assert row['customer']=='Schotpoort'
    assert row['customer_po']==row['barcode']==''
    assert row['item_no']=='934811.024'
    assert row['outer_pack']=='24'
    assert row['quantity']=='30000'
    assert row['unit_price_usd']=='1.094'
    assert row['amount_usd']=='32820.00'
    assert row['ship_date']==row['requested_delivery_date']=='2026-01-20'
    assert row['flags'][0]['code']=='goliath_delivery_date_conflict'
    assert row['flags'][0]['level']=='warning'
    assert '2025-01-20' in row['flags'][0]['text']


def test_matching_european_dates_have_no_conflict():
    assert parse_po(european_po(note_date='20-JAN-2026'),'EU.pdf')[0]['flags']==[]


def test_product_delivery_date_does_not_claim_inspection_fallback():
    row=parse_po(european_po(note_date='10-JAN-2026',date_label='INSPECTION DATE:'),'EU.pdf')[0]
    assert row['ship_date']=='2026-01-20'
    assert row['inspection_date']=='2026-01-10'
    assert row['flags']==[]


@pytest.mark.parametrize('value', ['14 SEPT 2026','14-SEPT-2026','14 September 2026','14 Sep 2026'])
def test_september_date_spellings(value):
    from app.services.huadeng_order_legacy.goliath_po_parser import _date
    assert _date(value)=='2026-09-14'


def test_invalid_calendar_date_is_not_guessed():
    from app.services.huadeng_order_legacy.goliath_po_parser import _date
    with pytest.raises(ValueError,match='日期无法识别'):
        _date('31 SEPT 2026')


def test_unrecognized_product_layout_still_rejected():
    with pytest.raises(ValueError,match='未识别商品表格'):
        parse_po(european_po(unknown_layout=True),'unknown.pdf')


@pytest.fixture
def preview_client(monkeypatch):
    # Exercise the actual route, parser, finalization and response model while
    # isolating auth/DB dependencies; no business database is started or changed.
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.api import customer_order as api
    app=FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[api.get_db]=lambda: None
    app.dependency_overrides[api.get_current_user]=lambda: object()
    monkeypatch.setattr(api,'_ensure_customer_order_permission',lambda *a: None)
    monkeypatch.setattr(api,'_has_duplicate_confirmation_permission',lambda *a: False)
    with TestClient(app) as client:
        yield client


@pytest.mark.parametrize('document', [po, european_po], ids=['far-east','european'])
def test_preview_endpoint_serializes_warnings_without_500(preview_client, document):
    response=preview_client.post('/api/customer-orders/goliath/preview-batch',
        data={'factory_id':'huadeng','received_date':'2026-09-11'},
        files=[('po_files',('customer.pdf',document(),'application/pdf')),
               ('schedule_file',('latest.xlsx',template(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))])
    assert response.status_code==200,response.text
    data=response.json()
    assert data['preview_fingerprint']
    assert data['summary']['total']==1
    assert data['summary']['blocked']==0
    for issue in data['rows'][0]['issues']:
        assert isinstance(issue['skip_key'],str)
        assert isinstance(issue['skip_label'],str)
    if document is european_po:
        assert data['rows'][0]['units_per_carton']=='24'
        assert data['rows'][0]['carton_count']=='1250'
        assert data['rows'][0]['requested_ship_date']=='2026-01-20'


def test_european_export_three_sheets_preserves_history():
    original=template()
    args=dict(factory_id='huadeng',customer_code='goliath',received_date='2026-09-11',
        po_files=[('EU.pdf',european_po())],schedule_file_name='latest.xlsx',schedule_content=original)
    output,_,preview=service.export_unified_customer_schedule(**args)
    row=preview['rows'][0]
    assert row['date_code']=='RR0096'
    book=openpyxl.load_workbook(BytesIO(output))
    item=book[service.ITEM_SHEET]
    assert (item['D5'].value,item['H5'].value,item['K5'].value,item['L5'].value)==('T_10130-PEU25_00642','934811.024',30000,24)
    assert item['Z5'].value.date().isoformat()=='2026-01-20'
    assert item['D4'].value=='HIST'
    for name in (service.ORDER_SHEET,service.REVIEW_SHEET):
        assert book[name]['P4'].value==1.094
        assert book[name]['C4'].value=='T_10130-PEU25_00642'
        assert book[name]['G4'].value=='934811.024'
    repeated=service.create_unified_customer_preview(**{**args,'schedule_content':output})
    assert any(i['code']=='duplicate_existing_order' for i in repeated['rows'][0]['issues'])
