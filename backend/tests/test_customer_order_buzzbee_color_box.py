"""Synthetic contracts for BuzzBee's separate unified-review packaging code."""
from io import BytesIO

import openpyxl
import pytest

from app.services import customer_order_buzzbee_color_box as boxes
from app.services import customer_order_huaxing_unified as mapping
from app.services import customer_order_unified as unified
from test_customer_order_huaxing_unified import template


CODE = "87654-04-28-TEST"
OTHER = "87654-05-28-TEST"


def review_order(book, number=4, so="SO-A", item="87654-1", customer="WMU", note=None):
    sheet = book["正单评审表"]
    for column, value in ((3, "CONTRACT-A"), (4, so), (6, customer), (7, item), (20, note)):
        sheet.cell(number, column).value = value


def item_order(book, number=4, so="SO-A", item="87654-1", customer="WMU", code=CODE, lane=0):
    sheet = book[mapping.BUZZBEE_SHEETS[lane]]
    for column, value in ((5, so), (7, customer), (8, item), (20, code)):
        sheet.cell(number, column).value = value


def snapshot(sheet):
    return {cell.coordinate: (cell.value, tuple(cell._style or [0] * 9))
            for row in sheet for cell in row if cell.value is not None or cell.has_style}


def roundtrip(book):
    buffer = BytesIO()
    book.save(buffer)
    return openpyxl.load_workbook(BytesIO(buffer.getvalue()))


def test_append_after_body_occupied_column_preserves_notes_and_all_original_cells():
    book = template()
    review = book["正单评审表"]
    review_order(book, so="", note=CODE)  # Actual historical layout has no SO.
    review["U15"] = '=IFERROR(VLOOKUP(A15,[1]Sheet1!A:B,2,FALSE),"")'
    # Compare saved representations: XLSX normalizes an empty literal to None.
    originals = {sheet.title: snapshot(sheet) for sheet in roundtrip(book)}
    column = boxes.prepare_review(book)
    assert column == 22 and review["V3"].value == "彩盒编号"
    assert review["V4"].value == CODE and review["T4"].value == CODE
    assert review["V15"].value is None
    reopened = roundtrip(book)
    for name, cells in originals.items():
        for coordinate, (value, style) in cells.items():
            assert reopened[name][coordinate].value == value
            assert tuple(reopened[name][coordinate]._style or [0] * 9) == style
    assert boxes.prepare_review(reopened) == column
    assert reopened["正单评审表"].max_column == column


def test_unique_full_identity_can_supply_history_code_and_notes_need_no_cross_row_lookup():
    book = template()
    item_order(book)
    review_order(book, note="保留业务备注")
    review_order(book, number=5, so="", item="UNRELATED", note=OTHER)
    col = boxes.prepare_review(book)
    assert book["正单评审表"].cell(4, col).value == CODE
    assert book["正单评审表"].cell(5, col).value == OTHER
    assert book["正单评审表"]["T4"].value == "保留业务备注"


@pytest.mark.parametrize("changes", [
    {"so": "SOA"}, {"item": "87654"}, {"customer": "OTHER"},
])
def test_history_never_relaxes_reference_variant_or_customer(changes):
    book = template()
    item_order(book, **changes)
    review_order(book)
    col = boxes.prepare_review(book)
    assert book["正单评审表"].cell(4, col).value is None


@pytest.mark.parametrize("note", [
    "彩盒", f"请改为 {CODE}", f"{CODE} / {OTHER}", f"{CODE}-{OTHER}", "=A1",
])
def test_unproven_or_formula_notes_do_not_become_codes(note):
    book = template()
    review_order(book, so="", note=note)
    col = boxes.prepare_review(book)
    assert book["正单评审表"].cell(4, col).value is None


@pytest.mark.parametrize("code", [f"{CODE}(r1)", f"{CODE}-HT", f"{CODE}-HT(r2)"])
def test_single_complete_code_keeps_all_suffixes_in_history_and_source_po(code):
    from app.services.customer_order_buzzbee import parse_po
    from test_customer_order_buzzbee import build_wmu_po

    book = template()
    review_order(book, so="", note=code)
    col = boxes.prepare_review(book)
    assert book["正单评审表"].cell(4, col).value == code
    po = openpyxl.load_workbook(BytesIO(build_wmu_po()))
    po["Packaging Box"]["B1"] = code
    buffer = BytesIO()
    po.save(buffer)
    parsed = parse_po("synthetic-po.xlsx", buffer.getvalue())
    assert len(parsed) == 7
    assert all(line.values["packaging"] == code for line in parsed)


@pytest.mark.parametrize("duplicate", [False, True])
def test_conflicting_note_or_multiple_item_matches_stay_empty(duplicate):
    book = template()
    item_order(book)
    review_order(book, note=OTHER if not duplicate else CODE)
    if duplicate:
        item_order(book, number=5, code=OTHER)
    col = boxes.prepare_review(book)
    assert book["正单评审表"].cell(4, col).value is None


@pytest.mark.parametrize("existing", [OTHER, '=""', " ", 0])
def test_existing_manual_or_formula_value_is_never_overwritten(existing):
    book = template()
    review_order(book, note=CODE)
    review = book["正单评审表"]
    review["W3"], review["W4"] = "彩盒编号", existing
    assert boxes.prepare_review(book) == 23
    assert review["W4"].value == existing and review.max_column == 23


def test_cancelled_and_shipped_sections_are_not_backfilled():
    book = template()
    review = book["正单评审表"]
    review_order(book, note=CODE)
    review["A5"] = "已走货订单"
    review_order(book, number=6, note=CODE)
    review_order(book, number=15, note=CODE)
    col = boxes.prepare_review(book)
    assert review.cell(4, col).value == CODE
    assert review.cell(6, col).value is None and review.cell(15, col).value is None


@pytest.mark.parametrize("layout", ["duplicate", "merged", "common"])
def test_invalid_existing_header_is_rejected_without_changing_workbook(layout):
    book = template()
    review = book["正单评审表"]
    review["U3"] = "彩盒编号"
    if layout == "duplicate":
        review["V3"] = " 彩盒编号 "
    elif layout == "merged":
        review.merge_cells("U3:V3")
    else:
        review["U3"], review["G3"] = None, "彩盒编号"
    original = snapshot(review)
    with pytest.raises(unified.CustomerOrderUnifiedError, match="彩盒编号"):
        boxes.prepare_review(book)
    assert snapshot(review) == original


def test_new_column_extends_existing_print_area_without_moving_titles():
    book = template()
    review = book["正单评审表"]
    review.print_area = "A1:T12"
    review.print_title_rows = "1:3"
    review.merge_cells("A1:T1")
    boxes.prepare_review(book)
    assert "$U$12" in review.print_area
    assert review.print_title_rows == "$1:$3"
    assert str(review.merged_cells) == "A1:T1"


@pytest.mark.parametrize("packaging", [CODE, f"{CODE}-HT(r2)", "", "普通彩盒"])
def test_new_rows_link_exact_item_lane_and_clear_any_inherited_formula(packaging):
    book = template()
    review = book["正单评审表"]
    review_order(book, note="原备注")
    review["U3"], review["U4"] = "彩盒编号", '="OLD-CODE"'
    lane = mapping.BUZZBEE_SHEETS[1]
    row = dict(id="new", reference_no="NEW-SO", contract_no="NEW-CONTRACT",
               product_no="87654-2", customer_name="WMU", quantity="12",
               units_per_carton="6", packaging=packaging, item_sheet_name=lane,
               issues=[], _huaxing_customer="buzzbee")
    mapping.export_rows(book, [row])
    ir = next(n for n in range(4, book[lane].max_row + 1) if book[lane].cell(n, 5).value == "NEW-SO")
    rr = next(n for n in range(4, review.max_row + 1) if review.cell(n, 4).value == "NEW-SO")
    assert review["U4"].value == '="OLD-CODE"'
    assert review.cell(rr, 21).value == (
        f'=IF(LEN(\'{lane}\'!T{ir})=0,"",\'{lane}\'!T{ir})'
        if packaging.startswith(CODE) else None
    )
    assert book["接单表"].max_column == 20
    assert book[lane].max_column == 30


def test_other_customers_do_not_add_buzzbee_header():
    book = template("dickie")
    mapping.export_rows(book, [dict(
        id="other", reference_no="OTHER", product_no="OTHER-ITEM", quantity="6",
        units_per_carton="6", item_sheet_name="ITEM表", issues=[], _huaxing_customer="dickie",
    )])
    assert not any(c.value == boxes.HEADER for c in book["正单评审表"][3])
