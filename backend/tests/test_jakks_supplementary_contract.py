from io import BytesIO
import openpyxl
import pytest

from app.services import customer_order_unified as unified
from app.services import customer_order_huadeng as huadeng
from app.services.huadeng_order_legacy import jakks_po_parser as parser
from app.services.huadeng_order_legacy.jakks_new_order_writer import records_from_order


# Product-table facts from the seven supplementary-contract sample layouts.
SAMPLES = [
    ('811784', '5475162', '48NSNBBP', '416984-PB', 'SONIC RINGS', 424, 4, '24-Nov-2026', 'POLYBAG'),
    ('811995', '5482226', '2551020438', '655094-PQ', 'MAUI WAVE DAZZLER JUMP ROPE', 2346, 6, '13-Nov-2026', 'BACKER CARD'),
    ('811996', '5482227', '8150973261', '655094-PQ', 'MAUI WAVE DAZZLER JUMP ROPE', 7182, 6, '25-Nov-2026', 'BACKER CARD'),
    ('812084', '5482433', '2551020439', '655044-V5', 'WAVE DAZZLER HOPPER VARIANT 4-PQ', 3112, 4, '13-Nov-2026', 'BACKCARD'),
    ('812085', '5482434', '8150973266', '655044-V5', 'WAVE DAZZLER HOPPER VARIANT 4-PQ', 7548, 4, '25-Nov-2026', 'BACKCARD'),
    ('811958', '67092', '236872', '10164J-V1', '8 IN BE MY BABY OLL W4 PDQ 1L-PQ-CVS', 3000, 12, '24-Nov-2026', 'OPEN TRAY BOX'),
    ('812341', '67115', '236894', '11180J', 'DP STC GFTABL GLMR CLLCT 3 S27 5L-GEN-PQ', 1696, 16, '20-Nov-2026', 'WINDOW BOX'),
]


def contract_text(sample=SAMPLES[0], *, formal=False):
    contract, confirmation, po, item, description, quantity, pack, ship, notes = sample
    prices = f' 2.0000 {quantity * 2:.2f}' if formal else ''
    return f'''{'CONTRACT' if formal else 'SUPPLEMENTARY CONTRACT'}
JAKKS PACIFIC (H.K. ) LIMITED
Contract #: {contract}
Printed Date: 2-Sep-2026 Order Date: 17-Aug-2026
Customer PO: {po}
Confirmation No: {confirmation} UNITED STATES
Ultimate Consignee Name: JAKKS PACIFIC, INC.
ITEM NUMBER DESCRIPTION QUANTITY PRICE EXTENDED PRICE
{item} {description} {quantity}{prices} USD
Customer Item No.: {item} {quantity // pack:,} CA
UPC: 199460009811 Outer Pack: {pack}
Stock #: {item}
Country of Origin: CHINA Line#: 1.000 Request Date: {ship}
NOTES: {notes}
CARTON MARKING
ORDER NOTES: REMARKS:
ITEM NO. TO BE SHOWN
{item}
Page 1 of 2
{'CONTRACT' if formal else 'SUPPLEMENTARY CONTRACT'}
JAKKS PACIFIC (H.K. ) LIMITED
Contract #: {contract}
TOTAL: USD NO DOLLARS AND ZERO CENTS
TOTAL USD
Page 2 of 2
'''


def workbook_bytes():
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name in unified.SHEETS:
        sheet = book.create_sheet(name)
        headers = unified.ITEM_HEADERS if name == unified.ITEM_SHEET else unified.ORDER_HEADERS
        for col, label in enumerate(headers, 1):
            sheet.cell(3, col, label)
        sheet.cell(6, 1, '取消单')
        sheet.cell(8, 1, '已走货订单')
        sheet.cell(9, 2, '保留历史资料')
    item = book[unified.ITEM_SHEET]
    for col, val in {4:'OLD', 5:'OLD', 7:'JAKKS PACIFIC, INC. UNITED STATES', 8:'416984-PB', 9:'中文品名', 11:100, 12:4}.items():
        item.cell(4, col, val)
    for name in (unified.ORDER_SHEET, unified.REVIEW_SHEET):
        sheet = book[name]
        for col, val in {3:'OLD', 4:'OLD', 7:'416984-PB', 14:99}.items():
            sheet.cell(4, col, val)
    data = BytesIO()
    book.save(data)
    return data.getvalue()


@pytest.mark.parametrize('sample', SAMPLES, ids=[s[0] for s in SAMPLES])
def test_unpriced_supplementary_contract(sample, monkeypatch, tmp_path):
    monkeypatch.setattr(parser, 'extract_text', lambda _: (contract_text(sample), False))
    parsed = parser.parse_po(tmp_path / 'customer-SUP.pdf')
    assert parsed['document_kind'] == 'supplementary_contract'
    assert (parsed['contract_no'], parsed['confirmation_no'], parsed['customer_po']) == sample[:3]
    assert len(parsed['lines']) == 1
    line = parsed['lines'][0]
    assert (line['item_no'], line['po_description'], line['quantity'], line['outer_pack']) == sample[3:7]
    assert line['cartons'] == sample[5] / sample[6]
    assert line['ship_date'] == parser._english_date(sample[7])
    assert line['product_packaging'] == sample[8]
    assert line['unit_price_usd'] is line['total_usd'] is None
    assert any('未提供单价' in w for w in parsed['warnings'])
    assert records_from_order(parsed)[0]['unit_price_usd'] == ''


def test_numeric_description_is_not_treated_as_price(monkeypatch, tmp_path):
    sample = list(SAMPLES[0])
    sample[4] = 'TOY SERIES 3 4'
    monkeypatch.setattr(parser, 'extract_text', lambda _: (contract_text(sample), False))
    line = parser.parse_po(tmp_path / 'supplement.pdf')['lines'][0]
    assert line['quantity'] == 424
    assert line['po_description'] == 'TOY SERIES 3 4'
    assert line['unit_price_usd'] is None


def test_priced_supplement_retains_explicit_prices(monkeypatch, tmp_path):
    text = contract_text(formal=True).replace('CONTRACT', 'SUPPLEMENTARY CONTRACT')
    monkeypatch.setattr(parser, 'extract_text', lambda _: (text, False))
    line = parser.parse_po(tmp_path / 'SUP.pdf')['lines'][0]
    assert (line['quantity'], line['unit_price_usd'], line['total_usd']) == (424, 2, 848)


@pytest.mark.parametrize('tail', ['424 2.0000 USD', '424 848.00 USD', '400 USD'])
def test_partial_price_or_inconsistent_packing_cannot_shift_quantity(tail, monkeypatch, tmp_path):
    text = contract_text().replace('424 USD', tail)
    monkeypatch.setattr(parser, 'extract_text', lambda _: (text, False))
    with pytest.raises(ValueError, match='数量与箱数'):
        parser.parse_po(tmp_path / 'SUP.pdf')


def test_ambiguous_partial_price_without_packing_is_blocked(monkeypatch, tmp_path):
    text = contract_text().replace('424 USD', '424 2.0000 USD').replace('106 CA', '').replace('Outer Pack: 4', '')
    monkeypatch.setattr(parser, 'extract_text', lambda _: (text, False))
    with pytest.raises(ValueError, match='存在歧义'):
        parser.parse_po(tmp_path / 'SUP.pdf')


def test_regular_po_still_requires_prices_and_cancelled_file_stays_blocked(monkeypatch, tmp_path):
    monkeypatch.setattr(parser, 'extract_text', lambda _: (contract_text().replace('SUPPLEMENTARY CONTRACT', 'CONTRACT'), False))
    with pytest.raises(ValueError, match='没有识别到产品明细'):
        parser.parse_po(tmp_path / 'formal.pdf')
    with pytest.raises(ValueError, match='取消单'):
        parser.parse_po(tmp_path / 'customer-CXL.pdf')
    monkeypatch.setattr(parser, 'extract_text', lambda _: (contract_text(formal=True), False))
    with pytest.raises(ValueError, match='正文未识别'):
        parser.parse_po(tmp_path / 'customer-SUP.pdf')


def test_supplement_exports_and_followup_po_does_not_duplicate(monkeypatch):
    monkeypatch.setattr(parser, 'extract_text', lambda p: (contract_text(formal='formal' in p.name), False))
    args = dict(factory_id='huadeng', customer_code='jakks', received_date='2026-09-12', schedule_file_name='排期.xlsx')
    original = workbook_bytes()
    output, _, preview = unified.export_unified_customer_schedule(
        **args, po_files=[('order-SUP.pdf', b'supplement')], schedule_content=original,
    )
    assert preview['summary']['total'] == 1
    assert preview['summary']['blocked'] == 0
    row = preview['rows'][0]
    assert row['order_type'] == '补充合同'
    assert row['unit_price_hkd'] == row['amount_hkd'] == ''
    assert any(i['code'] == 'supplementary_contract' for i in row['issues'])
    book = openpyxl.load_workbook(BytesIO(output))
    item = book[unified.ITEM_SHEET]
    assert [item.cell(5, c).value for c in (2, 4, 5, 6, 8, 11, 12, 20)] == [
        '补充合同', '811784', '5475162', '48NSNBBP', '416984-PB', 424, 4, 'POLYBAG',
    ]
    assert item.cell(5, 26).value.date().isoformat() == '2026-11-24'
    for name in unified.SHEETS:
        assert book[name].cell(9, 2).value == '保留历史资料'
    for name in (unified.ORDER_SHEET, unified.REVIEW_SHEET):
        assert [book[name].cell(5,c).value for c in range(3,8)] == [item.cell(5,c).value for c in range(4,9)]
        assert book[name].cell(4, 14).value == 99
    later = unified.create_unified_customer_preview(
        **args, po_files=[('formal.pdf', b'formal')], schedule_content=output,
    )
    assert later['summary']['blocked'] == 1
    assert any(i['code'] == 'duplicate_existing_order' for i in later['rows'][0]['issues'])
    with pytest.raises(unified.CustomerOrderUnifiedError):
        unified.export_unified_customer_schedule(**args, po_files=[('formal.pdf', b'formal')], schedule_content=output)
    batch = unified.create_unified_customer_preview(
        **args, po_files=[('order-SUP.pdf', b'supplement'), ('formal.pdf', b'formal')], schedule_content=original,
    )
    assert batch['summary']['total'] == 1


def test_legacy_batch_accepts_valid_sup_filename(monkeypatch):
    monkeypatch.setattr(parser, 'extract_text', lambda _: (contract_text(), False))
    monkeypatch.setattr(huadeng.jakks_schedule, 'load_dataset', lambda _: {'records': []})
    batch = huadeng._prepare_jakks([('order-SUP.pdf', b'pdf'), ('order-CXL.pdf', b'pdf')], 'schedule.xlsx', b'xlsx')
    assert len(batch.records) == 1
    assert batch.records[0]['quantity'] == 424
    assert any('CXL' in w for w in batch.warnings)
