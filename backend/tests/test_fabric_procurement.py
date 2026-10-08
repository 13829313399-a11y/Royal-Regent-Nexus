import importlib
import json
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook
from sqlalchemy import select, func
from test_molding_sample_api import login_as, make_client

BASE = "/api/fabric-warehouse/procurement"
FIELDS = ["订单日期", "订单号", "供应商", "生产单号", "计划跟踪号", "物料编码", "物料名称", "订单数量", "基本单位", "单价", "合同号", "交货明细", "入库数量", "送货单号", "采购明细ID"]


def row(**changes):
    return {"订单日期": "2026/9/16", "订单号": "CGDD001", "供应商": "采购供应商", "生产单号": "P001", "计划跟踪号": "计划001",
            "物料编码": "000123", "物料名称": "白色布 58寸", "订单数量": 100, "基本单位": "码", "单价": 3.99,
            "合同号": "SC001", "交货明细": "20+30", "入库数量": 50, "送货单号": "0000123", "采购明细ID": "", **changes}


def book(pending=None, returned=None, reverse=False):
    wb = Workbook(); wb.remove(wb.active)
    for name, rows in [("未回物料", pending), ("已回料", returned)]:
        if rows is None:
            continue
        ws = wb.create_sheet(name)
        if name == "未回物料":
            ws.append(["说明文字，不能当作数据"])
        headers = list(reversed(FIELDS)) if reverse else FIELDS
        if name == "已回料":
            headers = [h.replace("生产单号", "放产单号") for h in headers]
        ws.append(headers)
        for data in rows:
            ws.append([data.get(h.replace("放产单号", "生产单号")) for h in headers])
    wb.create_sheet("留存历史").append(["不应导入"])
    out = BytesIO(); wb.save(out); return out.getvalue()


def upload(client, content, *, apply=False, **kwargs):
    data = {"factory_id": "huakang-c", "scope": "ALL", **kwargs}
    return client.post(BASE + ("/imports/apply" if apply else "/imports/preview"), data=data, files={"file": ("来源.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def commit(client, content, preview, request=None, **kwargs):
    return upload(client, content, apply=True, preview_token=preview["preview_token"], request_id=request or str(uuid4()), confirmed="true", acknowledge_excluded="true", **kwargs)


def stored(client, status="ALL", **kwargs):
    response = client.get(BASE + "/lines", params={"factory_id": "huakang-c", "status": status, **kwargs})
    assert response.status_code == 200, response.text
    return response.json()


def test_import_preview_history_partial_receipts_and_retry_preserve_evidence_without_stock(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        content = book([row()], [row(订单号="CGDD002", 入库数量=100, 交货明细="100")])
        preview = upload(client, content)
        assert preview.status_code == 200, preview.text
        data = preview.json(); assert data["counts"]["new"] == 2
        assert stored(client)["total"] == 0
        assert data["items"][0]["facts"]["reported_received_quantity"] == "50"
        assert data["items"][0]["facts"]["material_code"] == "000123"
        assert data["items"][0]["facts"]["delivery_note_no"] == "0000123"
        assert "留存历史" in data["ignored_sheets"]
        assert "raw" not in data["items"][0]
        request = str(uuid4())
        result = commit(client, content, data, request)
        assert result.status_code == 200, result.text
        assert result.json()["stock_posted"] is False
        repeated = commit(client, content, data, request)
        assert repeated.json() == result.json()
        assert stored(client, "PENDING")["total"] == 1
        assert stored(client, "RETURNED")["total"] == 1
        line = stored(client, search="000123")["items"][0]
        detail = client.get(BASE + "/lines/" + line["id"], params={"factory_id": "huakang-c"}).json()
        assert detail["evidence"][0]["source_name"] == "来源.xlsx"
        assert detail["evidence"][0]["raw"]["values"]
        assert upload(client, content).json()["counts"]["unchanged"] == 2
        altered = book([row(订单数量=101)])
        assert commit(client, altered, data, request).status_code == 409
        assert stored(client)["total"] == 2
        dbm = importlib.import_module("app.db")
        model = importlib.import_module("app.models.carton_procurement")
        with dbm.SessionLocal() as db:
            assert db.scalar(select(func.count()).select_from(model.CartonInventoryMovement)) == 0


def test_same_material_split_production_and_multiline_reorder_is_not_collapsed(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        values = [row(订单数量=10), row(订单数量=20), row(生产单号="P002", 订单数量=30)]
        content = book(values); data = upload(client, content).json()
        assert data["counts"]["new"] == 3
        assert commit(client, content, data).status_code == 200
        original_ids = {r["id"] for r in stored(client)["items"]}
        reordered = upload(client, book(list(reversed(values)), reverse=True)).json()
        assert reordered["counts"]["unchanged"] == 3
        assert original_ids == {r["id"] for r in stored(client)["items"]}
        changed = book([row(订单数量=11), values[1], row(生产单号="P003")])
        updated = upload(client, changed).json()
        assert updated["counts"]["blocked"] == 2
        assert updated["counts"]["new"] == 1
        rejected = upload(client, changed, apply=True, preview_token=updated["preview_token"], request_id=str(uuid4()), confirmed="true")
        assert rejected.status_code == 422
        assert commit(client, changed, updated).status_code == 200
        assert stored(client)["total"] == 4


def test_update_is_previewed_stale_file_or_global_revision_cannot_commit(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        content = book([row()]); data = upload(client, content).json()
        assert commit(client, content, data).status_code == 200
        changed = book([row(订单数量=120)])
        new = upload(client, changed).json()
        assert new["counts"]["updated"] == 1
        assert "ordered_quantity" in new["items"][0]["changes"]
        assert commit(client, changed, new).status_code == 200
        assert commit(client, changed, new).status_code == 409
        detail = client.get(BASE + "/lines/" + stored(client)["items"][0]["id"], params={"factory_id": "huakang-c"}).json()
        assert len(detail["evidence"]) == 2
        assert detail["evidence"][0]["before"]["ordered_quantity"] == "100"
        assert detail["evidence"][0]["after"]["ordered_quantity"] == "120"


def test_unknown_formula_zero_price_units_and_historical_columns_are_not_guessed(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        content = book([row(交货明细="", 入库数量="#VALUE!", 单价=0), row(生产单号="P002", 交货明细="=EVALUATE(A1)", 入库数量="#VALUE!", 基本单位="千克(转算专用)")],
                       [row(订单号="CGDD003", 基本单位="140")])
        data = upload(client, content).json()
        assert data["counts"]["blocked"] == 1
        assert data["items"][0]["facts"]["reported_received_quantity"] is None
        assert data["items"][0]["facts"]["unit_price"] == "0"
        assert data["items"][1]["facts"]["unit"] == "千克(转算专用)"
        assert data["items"][1]["facts"]["reported_received_quantity"] is None
        assert commit(client, content, data).status_code == 200
        assert stored(client)["total"] == 2


def test_auth_factory_department_denies_preview_scope_and_no_auth_fail_closed(monkeypatch):
    with make_client(monkeypatch) as client:
        content = book([row()])
        assert upload(client, content).status_code == 401
        login_as(client, "admin")
        assert upload(client, content, factory_id="huakang-a").status_code == 422
        assert upload(client, content, factory_id="group").status_code == 422
        assert client.get(BASE + "/lines").status_code == 422
        preview = upload(client, content, scope="PENDING").json()
        assert commit(client, content, preview, scope="ALL").status_code == 409
        login_as(client, "engineer")
        assert upload(client, content).status_code == 403
        assert client.get(BASE + "/lines", params={"factory_id": "huakang-c"}).status_code == 403


def test_parser_duplicate_stable_ids_and_invalid_headers_are_blocked():
    parser = importlib.import_module("app.services.fabric_procurement_parser")
    data = parser.parse_workbook("example.xlsx", book([row()]), "PENDING")
    assert data["rows"][0]["facts"]["order_date"] == "2026-09-16"
    with pytest.raises(Exception) as invalid:
        parser.parse_workbook("example.xlsx", b"not a workbook")
    assert invalid.value.status_code == 422


def test_field_lengths_match_database_and_direct_source_contract():
    parser = importlib.import_module("app.services.fabric_procurement_parser")
    for field in ("source_line_id", "order_no", "material_code", "old_material_code"):
        valid = {"order_no": "PO", "supplier": "Supplier", "material_code": "MAT", "material_name": "Fabric", "unit": "码", "ordered_quantity": "10", field: "X" * 128}
        assert not parser.normalize_row(valid, "PENDING")[2]
        valid[field] += "X"
        assert f"字段过长：{field}" in parser.normalize_row(valid, "PENDING")[2]
    valid = {"order_no": "PO", "supplier": "Supplier", "material_code": "MAT", "material_name": "Fabric", "unit": "码", "ordered_quantity": "10", "source_line_id": "#REF!"}
    assert parser.normalize_row(valid, "PENDING")[2]


def test_numeric_identifier_zero_format_is_preserved_and_complex_format_is_blocked():
    parser = importlib.import_module("app.services.fabric_procurement_parser")
    wb = Workbook(); ws = wb.active; ws.title = "未回物料"
    ws.append(FIELDS); ws.append([row().get(h) for h in FIELDS])
    cell = ws.cell(2, FIELDS.index("物料编码") + 1); cell.value = 123; cell.number_format = "000000"
    out = BytesIO(); wb.save(out)
    parsed = parser.parse_workbook("source.xlsx", out.getvalue())
    assert parsed["rows"][0]["facts"]["material_code"] == "000123"
    assert not parsed["rows"][0]["errors"]
    assert parsed["rows"][0]["raw"]["number_formats"][cell.coordinate]["stored"] == "123"
    cell.number_format = "000-000"; out = BytesIO(); wb.save(out)
    assert parser.parse_workbook("source.xlsx", out.getvalue())["rows"][0]["errors"]
    cell.value = 123.5; cell.number_format = "0"; out = BytesIO(); wb.save(out)
    assert parser.parse_workbook("source.xlsx", out.getvalue())["rows"][0]["errors"]
    for field in ("supplier", "production_no", "material_name"):
        valid = {"order_no": "PO", "supplier": "Supplier", "material_code": "MAT", "material_name": "Fabric", "unit": "码", "ordered_quantity": "10", field: "X" * 255}
        assert not parser.normalize_row(valid, "PENDING")[2]
        valid[field] += "X"
        assert f"字段过长：{field}" in parser.normalize_row(valid, "PENDING")[2]


def test_malformed_structure_and_empty_row_amplification_are_rejected(monkeypatch):
    from zipfile import ZipFile, ZIP_DEFLATED
    parser = importlib.import_module("app.services.fabric_procurement_parser")
    out = BytesIO()
    with ZipFile(out, "w", ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>')
        archive.writestr("xl/_rels/workbook.xml.rels", '<Relationships/>')
    with pytest.raises(Exception) as missing:
        parser.parse_workbook("source.xlsx", out.getvalue())
    assert missing.value.status_code == 422
    monkeypatch.setattr(parser, "MAX_SCANNED_ROWS", 2)
    with pytest.raises(Exception) as excessive:
        parser.parse_workbook("source.xlsx", book([row()]))
    assert excessive.value.status_code == 422


def test_mixed_document_ids_and_competing_alias_updates_never_duplicate_or_overwrite(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        mixed = book([row(采购明细ID="K1"), row(), row(订单号="CGDD002")])
        preview = upload(client, mixed).json()
        assert preview["counts"]["blocked"] == 2
        assert preview["counts"]["new"] == 1
        assert commit(client, mixed, preview).status_code == 200
        assert stored(client)["total"] == 1
        original = book([row()]); preview = upload(client, original).json()
        assert commit(client, original, preview).status_code == 200
        linked = book([row(采购明细ID="K1")])
        # Linking an upstream ID requires an explicit warehouse line binding.
        assert upload(client, linked).json()["counts"]["blocked"] == 1
        dbm = importlib.import_module("app.db")
        models = importlib.import_module("app.models.fabric_procurement")
        with dbm.SessionLocal() as db:
            line = db.scalar(select(models.FabricProcurementLine).where(models.FabricProcurementLine.order_no == "CGDD001"))
            facts = json.loads(line.payload_json); facts["source_line_id"] = "K1"
            line.payload_json = json.dumps(facts); db.commit()
        competing = book([row(采购明细ID="K1", 生产单号="P002", 订单数量=120), row(订单数量=130), row(订单号="CGDD003")])
        preview = upload(client, competing).json()
        assert preview["counts"]["blocked"] == 2
        assert preview["counts"]["updated"] == 0
        assert commit(client, competing, preview).status_code == 200
        persisted = stored(client, search="CGDD001")["items"]
        assert len(persisted) == 1
        assert persisted[0]["facts"]["ordered_quantity"] == "100"


def test_reserved_purchase_feed_reuses_ingestion_and_requires_explicit_legacy_links(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        content = book([row()]); preview = upload(client, content).json()
        assert commit(client, content, preview).status_code == 200
        legacy = stored(client)["items"][0]
        data = {"factory_id": "huakang-c", "feed": {"source_system": "KINGDEE", "snapshot_id": "SYNC001", "lines": [{
            "source_line_id": "KD-DETAIL-001", "status": "PENDING", "order_no": "CGDD001", "supplier": "采购供应商", "production_no": "P001", "plan_no": "计划001",
            "material_code": "000123", "material_name": "白色布 58寸", "unit": "码", "contract_no": "SC001", "ordered_quantity": "100", "reported_received_quantity": "50", "unit_price": "3.99", "order_date": "2026-09-16"}]}}
        unbound = client.post(BASE + "/sources/preview", json=data)
        assert unbound.status_code == 200, unbound.text
        assert unbound.json()["counts"]["blocked"] == 1
        data["feed"]["lines"][0].update(warehouse_line_id=legacy["id"], expected_line_revision=legacy["revision"])
        bound = client.post(BASE + "/sources/preview", json=data)
        assert bound.status_code == 200, bound.text
        assert bound.json()["counts"]["updated"] == 1
        confirm = {**data, "preview_token": bound.json()["preview_token"], "request_id": str(uuid4()), "confirmed": True}
        saved = client.post(BASE + "/sources/apply", json=confirm)
        assert saved.status_code == 200, saved.text
        assert stored(client)["total"] == 1
        assert stored(client)["items"][0]["id"] == legacy["id"]
        assert stored(client)["items"][0]["facts"]["source_line_id"] == "KD-DETAIL-001"
        data["feed"]["lines"][0].update(warehouse_line_id="", expected_line_revision=None, ordered_quantity="120")
        changed = client.post(BASE + "/sources/preview", json=data).json()
        assert changed["counts"]["updated"] == 1
        assert client.post(BASE + "/sources/apply", json={**data, "preview_token": changed["preview_token"], "request_id": str(uuid4()), "confirmed": True}).status_code == 200
        assert stored(client)["total"] == 1
        data["feed"]["lines"].append(data["feed"]["lines"][0].copy())
        duplicate = client.post(BASE + "/sources/preview", json=data).json()
        assert duplicate["counts"]["blocked"] == 2


def test_new_domain_checks_canonical_factory_department_and_explicit_denies_even_in_legacy(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "admin")
        from test_raw_material_api import create_fixed_position_user, login_fixed_position_user
        create_fixed_position_user("fabric_c", "position_warehouse_keeper", "pmc-warehouse", factory_id="huakang-c")
        create_fixed_position_user("fabric_d", "position_warehouse_keeper", "pmc-warehouse", factory_id="huakang-d")
        login_fixed_position_user(client, "fabric_c")
        assert stored(client)["total"] == 0
        assert upload(client, book([row()])).status_code == 200
        login_fixed_position_user(client, "fabric_d")
        assert client.get(BASE + "/lines", params={"factory_id": "huakang-c"}).status_code == 403
        auth = importlib.import_module("app.services.auth")
        api = importlib.import_module("app.api.fabric_procurement")
        dbm = importlib.import_module("app.db")
        models = importlib.import_module("app.models.auth")
        from dataclasses import replace
        with dbm.SessionLocal() as db:
            actor = auth.build_auth_context(db, db.get(models.AuthUser, "user-fabric_c"))
            denied = replace(actor, overrides=(auth.AuthOverrideContext(id="deny-fabric", permission_code="fabric_warehouse:import", effect="deny", factory_id="huakang-c", department="pmc-warehouse"),))
            with pytest.raises(Exception) as refused:
                api.authorize(db, denied, "huakang-c", "import")
            assert refused.value.status_code == 403
            misplaced = replace(actor, grants=tuple(replace(grant, department="engineering", unrestricted_department=True) for grant in actor.grants))
            for code in ("fabric_warehouse:read", "fabric_warehouse:import"):
                assert not auth.can(misplaced, code, "huakang-c", "pmc-warehouse")


def test_migration_upgrade_and_downgrade_protect_source_evidence(tmp_path):
    import sqlalchemy as sa
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from importlib.util import spec_from_file_location, module_from_spec
    from pathlib import Path
    path = Path(__file__).parents[1] / "alembic/versions/20261005_0134_fabric_procurement.py"
    spec = spec_from_file_location("fabric_migration", path); module = module_from_spec(spec); spec.loader.exec_module(module)
    engine = sa.create_engine("sqlite:///" + str(tmp_path / "migration.db"))
    with engine.begin() as connection:
        context = MigrationContext.configure(connection)
        with Operations.context(context):
            module.upgrade()
            assert "fabric_procurement_evidence" in sa.inspect(connection).get_table_names()
            connection.execute(sa.text("INSERT INTO fabric_procurement_lines VALUES ('L', 'huakang-c', 'KEY', 1, 'PENDING', 'PO', 'Supplier', 'MAT', 'PROD', '{}', 1, 'NOW')"))
            with pytest.raises(RuntimeError, match="证据"):
                module.downgrade()
            connection.execute(sa.text("DELETE FROM fabric_procurement_lines"))
            module.downgrade()
            assert "fabric_procurement_lines" not in sa.inspect(connection).get_table_names()


@pytest.mark.parametrize("same_request", [True, False])
def test_concurrent_saves_are_idempotent_or_reject_stale_preview(tmp_path, same_request):
    from concurrent.futures import ThreadPoolExecutor
    from copy import deepcopy
    from threading import Barrier
    from types import SimpleNamespace
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from fastapi import HTTPException
    service = importlib.import_module("app.services.fabric_procurement")
    parser = importlib.import_module("app.services.fabric_procurement_parser")
    models = importlib.import_module("app.models.fabric_procurement")
    engine = create_engine("sqlite:///" + str(tmp_path / "concurrent.db"), connect_args={"check_same_thread": False})
    tables = [m.__table__ for m in (models.FabricProcurementState, models.FabricProcurementLine, models.FabricProcurementImport, models.FabricProcurementEvidence)]
    importlib.import_module("app.db").Base.metadata.create_all(engine, tables=tables)
    parsed = parser.parse_workbook("source.xlsx", book([row()]))
    actor = SimpleNamespace(id="actor", display_name="Warehouse")
    with Session(engine) as db:
        token = service.preview(db, deepcopy(parsed), actor.id)["preview_token"]
    requests = [str(uuid4()), str(uuid4())]
    if same_request:
        requests[1] = requests[0]
    barrier = Barrier(2)
    def save(request):
        with Session(engine) as db:
            barrier.wait(timeout=10)
            try:
                return service.apply(db, deepcopy(parsed), actor, "source.xlsx", request, token, confirmed=True, acknowledge_excluded=True)
            except HTTPException as exc:
                return exc.status_code
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(save, requests))
    if same_request:
        assert results[0] == results[1] and isinstance(results[0], dict)
    else:
        assert sum(result == 409 for result in results) == 1
        assert sum(isinstance(result, dict) for result in results) == 1
    with Session(engine) as db:
        assert db.scalar(select(func.count()).select_from(models.FabricProcurementLine)) == 1
        assert db.scalar(select(func.count()).select_from(models.FabricProcurementImport)) == 1
