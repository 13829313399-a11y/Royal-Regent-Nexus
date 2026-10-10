import io
import json
import runpy
from pathlib import Path

import pytest
import sqlalchemy as sa
from alembic.migration import MigrationContext
from alembic.operations import Operations
from fastapi import HTTPException
from openpyxl import Workbook
from pypdf import PdfReader
from test_molding_sample_api import make_client
from test_carton_supplier_mark_pdf import prepare_layout, BASE
from test_carton_supplier_portal import revoke_supplier_permission, restore_supplier_permission
from test_carton_mark_api import _pdf_bytes

LAYOUTS = "/api/carton-supplier/carton-mark/layouts"


def workbook(**changes):
    book = Workbook(); sheet = book.active
    values = dict(item="00123", vendor="ZURU INC", product="UNICORN", content="24 PCS / CTN",
        gtin="00193052021172", vendor_item="00123", dimensions="12 x 18 x 20 CM", gross_weight="2.5 KG")
    values.update(changes)
    labels = ["Item", "Vendor name", "Product name", "Content", "GTIN", "Vendor product number", "Width/Length/Height", "Grossweight"]
    for row, (key, name) in enumerate(zip(values, labels), 1):
        sheet.cell(row, 1, name + ":"); sheet.cell(row, 2, values[key])
    stream = io.BytesIO(); book.save(stream); book.close(); return stream.getvalue()


def original(config=None, content=None):
    from app.schemas.carton_supplier_portal import MarkLayoutConfig
    return dict(excel_file_name="customer.xlsx", excel_bytes=content or workbook(), item_no="00123", contract_no="4500222793",
        layout=dict(name="Customer", version=1, reference_bytes=_pdf_bytes(), config=MarkLayoutConfig(**(config or {})).model_dump()))


def test_renderer_preserves_fields_leading_zero_and_generates_customer_modes():
    from app.services.carton_mark_layout_renderer import render
    for mode, pages in (("front_side", 1), ("front_only", 1), ("separate_pages", 2)):
        content, count, warnings = render(original(dict(mode=mode, side_address="Customer address\ncare@example.com")))
        reader = PdfReader(io.BytesIO(content))
        assert len(reader.pages) == count == pages
        text = "\n".join(page.extract_text() for page in reader.pages)
        assert "00193052021172" in text and "00123" in text and "UNICORN" in text
        assert "care@example.com" in text if mode != "front_only" else "care@example.com" not in text
        assert not any("缺少内容" in warning for warning in warnings)


def test_original_frame_has_red_vector_dividers_fold_and_black_content_in_every_mode():
    from app.services.carton_mark_layout_renderer import render
    from pypdf.generic import ContentStream
    for mode, expected in (("front_side", 1), ("front_only", 1), ("separate_pages", 2)):
        content, pages, _ = render(original(dict(mode=mode, frame_style="original_red", instructions="Confirm artwork before printing", side_address="Customer address\ncare@example.com")))
        reader = PdfReader(io.BytesIO(content)); assert pages == expected
        for page in reader.pages:
            operations = ContentStream(page.get_contents(), reader).operations
            assert any(op == b"RG" and tuple(round(float(number), 2) for number in args) == (.94, .23, .14) for args, op in operations)
            assert "00193052021172" in page.extract_text() and "UNICORN" in page.extract_text()
        # The folding flap begins inside the printable page at 2 mm.
        assert any(op == b"l" and abs(float(args[0]) - 2 * 72 / 25.4) < .01
            for args, op in ContentStream(reader.pages[0].get_contents(), reader).operations)
    content, _, _ = render(original())
    reader = PdfReader(io.BytesIO(content))
    assert not any(op == b"RG" and tuple(round(float(number), 2) for number in args) == (.94, .23, .14)
        for args, op in ContentStream(reader.pages[0].get_contents(), reader).operations)


def test_renderer_requires_unambiguous_fields_and_valid_barcode_and_never_clips_long_values():
    from app.services.carton_mark_layout_renderer import render, extract_fields
    book = Workbook(); sheet = book.active
    for row, values in enumerate((("Item:", "00123"), ("Product name:", "UNICORN"), ("GTIN:", "001"), ("GTIN:", "002")), 1):
        for column, value in enumerate(values, 1): sheet.cell(row, column, value)
    stream = io.BytesIO(); book.save(stream); book.close()
    with pytest.raises(HTTPException) as error:
        extract_fields("test.xlsx", stream.getvalue(), {})
    assert "多个不同内容" in error.value.detail
    assert extract_fields("test.xlsx", stream.getvalue(), {"gtin": "B3"})[0][1]["gtin"] == "001"
    with pytest.raises(HTTPException) as error: render(original(dict(barcode="ITF14"), workbook(gtin="123")))
    assert error.value.status_code == 422
    with pytest.raises(HTTPException) as error: render(original(content=workbook(product="Very long product " * 100)))
    assert "内容过长" in error.value.detail
    _, _, warnings = render(original(content=workbook(dimensions="", gross_weight="")))
    assert any("尺寸、毛重" in warning for warning in warnings)


def test_customer_versions_keep_reference_and_reject_stale_or_foreign_context(monkeypatch):
    with make_client(monkeypatch) as client:
        _, payload, _ = prepare_layout(client)
        scope = {key: payload[key] for key in ("factory_id", "order_id", "issue_id")}
        listed = client.get(LAYOUTS, params=scope)
        assert listed.status_code == 200 and listed.json()[0]["version"] == 1
        assert listed.json()[0]["config"]["frame_style"] == "plain"
        preview = client.get(f"{LAYOUTS}/{payload['layout_id']}/preview", params=scope)
        assert preview.status_code == 200 and preview.json()["preview_data_url"].startswith("data:image/png;base64,")
        update = {**scope, "name": "Customer revised", "config": json.dumps(dict(mode="front_only", frame_style="original_red")), "expected_version": 1, "base_template_id": payload["layout_id"]}
        saved = client.post(LAYOUTS, data=update)
        assert saved.status_code == 201, saved.text
        assert saved.json()["version"] == 2
        assert saved.json()["config"]["frame_style"] == "original_red"
        assert client.post(LAYOUTS, data=update).status_code == 409
        assert client.post(BASE, json=payload).status_code == 409
        assert client.get(LAYOUTS, params={**scope, "factory_id": "huakang-b"}).status_code == 403
        assert client.get(f"{LAYOUTS}/{payload['layout_id']}/preview", params={**scope, "order_id": "foreign"}).status_code == 404
        from app.db import SessionLocal
        from app.models.carton_mark import CartonMarkLayout
        with SessionLocal() as db:
            versions = list(db.scalars(sa.select(CartonMarkLayout).order_by(CartonMarkLayout.version)))
            assert [bytes(row.reference_content) for row in versions] == [_pdf_bytes(), _pdf_bytes()]
        revoke_supplier_permission("carton_supplier:edit")
        assert client.post(LAYOUTS, data={**update, "expected_version": 2, "base_template_id": saved.json()["id"]}).status_code == 403
        restore_supplier_permission("carton_supplier:edit")


def test_template_revalidates_permissions_after_reference_preparation(monkeypatch):
    with make_client(monkeypatch) as client:
        _, payload, _ = prepare_layout(client)
        from app.services import carton_supplier_mark_layout as service
        from app.db import SessionLocal
        from app.models.carton_mark import CartonMarkLayout
        validate = service.validate_reference
        def changed(*args):
            result = validate(*args)
            revoke_supplier_permission("carton_supplier:edit")
            return result
        monkeypatch.setattr(service, "validate_reference", changed)
        response = client.post(LAYOUTS, data={**{key: payload[key] for key in ("factory_id", "order_id", "issue_id")},
            "name": "revised", "config": "{}", "expected_version": 1, "base_template_id": payload["layout_id"]})
        assert response.status_code == 403, response.text
        with SessionLocal() as db: assert db.scalar(sa.select(sa.func.count()).select_from(CartonMarkLayout)) == 1


def test_reference_rejects_outside_regions_and_damaged_pdf():
    from app.services.carton_supplier_mark_layout import validate_reference
    from app.schemas.carton_supplier_portal import MarkLayoutConfig
    for content, config in ((b"invalid", {}), (_pdf_bytes(), dict(logo_region=dict(x=999, y=0, width=10, height=10)))):
        with pytest.raises(HTTPException) as error: validate_reference(content, MarkLayoutConfig(**config))
        assert error.value.status_code == 422


def test_generation_rejects_template_changed_while_rendering(monkeypatch):
    with make_client(monkeypatch) as client:
        _, payload, _ = prepare_layout(client)
        from app.services import carton_supplier_mark_pdf as service
        from app.db import SessionLocal
        from app.models.carton_mark import CartonMarkLayout, CartonMarkAsset
        def changed(original):
            with SessionLocal() as db:
                previous = db.get(CartonMarkLayout, payload["layout_id"])
                values = {column.name: getattr(previous, column.name) for column in CartonMarkLayout.__table__.columns}
                values.update(id="new-layout", version=2)
                db.add(CartonMarkLayout(**values)); db.commit()
            return _pdf_bytes(), 1, []
        monkeypatch.setattr(service, "prepare_pdf", changed)
        response = client.post(BASE, json=payload)
        assert response.status_code == 409 and "模板已更新" in response.json()["detail"]
        with SessionLocal() as db: assert db.scalar(sa.select(sa.func.count()).select_from(CartonMarkAsset)) == 1


def test_fixed_artwork_crop_renders_original_pixels_without_historical_text():
    from app.services.carton_supplier_mark_layout import crop_reference
    from PIL import Image
    content = crop_reference(_pdf_bytes(), 0, dict(x=1, y=1, width=10, height=12))
    with Image.open(io.BytesIO(content)) as picture:
        assert 83 <= picture.width <= 86 and 100 <= picture.height <= 104


def test_multiple_item_sheets_only_render_the_selected_order_item():
    from app.services.carton_mark_layout_renderer import render
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(workbook()))
    another = book.copy_worksheet(book.active); another.title = "Another item"
    another["B1"] = "FOREIGN-ITEM"; another["B3"] = "FOREIGN-PRODUCT"
    stream = io.BytesIO(); book.save(stream); book.close()
    data = original(content=stream.getvalue())
    content, pages, warnings = render(data)
    assert pages == 1 and any("跳过 1 张" in warning for warning in warnings)
    assert "FOREIGN-PRODUCT" not in PdfReader(io.BytesIO(content)).pages[0].extract_text()
    # A LOGO on a skipped item sheet must not supply artwork for this order.
    from openpyxl.drawing.image import Image as ExcelImage
    from PIL import Image
    book = load_workbook(io.BytesIO(stream.getvalue()))
    picture = io.BytesIO(); Image.new("RGB", (10, 10), "red").save(picture, "PNG"); picture.seek(0)
    book["Another item"].add_image(ExcelImage(picture), "C1")
    with_logo = io.BytesIO(); book.save(with_logo); book.close()
    _, _, warnings = render(original(content=with_logo.getvalue()))
    assert any("未找到唯一 LOGO" in warning for warning in warnings)
    data["item_no"] = "UNMATCHED"
    with pytest.raises(HTTPException) as error: render(data)
    assert "没有工作表匹配" in error.value.detail


def test_layout_migration_preserves_rows_factory_boundary_and_retained_versions(tmp_path, monkeypatch):
    migration = runpy.run_path(str(Path(__file__).parents[1] / "alembic/versions/20261010_0154_carton_mark_customer_layouts.py"))
    engine = sa.create_engine(f"sqlite:///{tmp_path / 'layouts.db'}")
    with engine.begin() as connection:
        connection.exec_driver_sql("PRAGMA foreign_keys=ON")
        connection.exec_driver_sql("CREATE TABLE carton_suppliers (id VARCHAR(96), factory_id VARCHAR(64), UNIQUE(id, factory_id))")
        connection.exec_driver_sql("INSERT INTO carton_suppliers VALUES ('S1', 'huaxing')")
        operations = Operations(MigrationContext.configure(connection))
        monkeypatch.setattr(migration["upgrade"].__globals__["op"], "get_bind", lambda: connection)
        monkeypatch.setattr(migration["upgrade"].__globals__["op"], "create_table", operations.create_table)
        migration["upgrade"]()
        assert connection.exec_driver_sql("SELECT * FROM carton_suppliers").all() == [("S1", "huaxing")]
        table = sa.Table("carton_mark_layouts", sa.MetaData(), autoload_with=connection)
        values = dict(id="L1", factory_id="huaxing", supplier_id="S1", customer_key="name:customer", customer_name="customer",
            name="fixed", version=1, config_json="{}", reference_name="old.pdf", reference_sha256="sha", reference_size=1,
            reference_content=b"x", created_by="U1", created_by_name="actor", created_at="today")
        connection.execute(table.insert().values(**values))
        for changes in (dict(id="L2"), dict(id="L2", factory_id="huakang-b")):
            with pytest.raises(sa.exc.IntegrityError): connection.execute(table.insert().values(**{**values, **changes}))
        with pytest.raises(RuntimeError, match="history is retained"): migration["downgrade"]()
