from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
from functools import lru_cache
import json

from fastapi import HTTPException
import openpyxl
import pytest
from sqlalchemy import func, select

from test_customer_order_ledger import engine, db, client, user, ship_body
from app.api import customer_order_ledger as api
from app.models.customer_order_ledger import OrderLedgerLine as Line, OrderLedgerSource as Source, OrderLedgerVersion as Version, OrderLedgerShipment as Shipment, OrderLedgerDispatch as Dispatch
from app.schemas.customer_order_ledger import HistorySelection, HistoryOpeningIn, ReasonIn, RestoreIn, DispatchIn
from app.services import customer_order_history as history, customer_order_ledger as ledger
from app.services.auth import get_current_user


@lru_cache(maxsize=1)
def workbook():
    w = openpyxl.Workbook()
    s = w.active
    s.title = 'ITEM表'
    headers = ['客出单日期', 'Contract No.', 'SO#/Reference', 'P/O#:', '客名', '产品编号', '數量', '客要求走货期', '备注', '箱唛', '客户专属字段']
    s.append([]); s.append([]); s.append(headers)
    s.append([date(2025, 1, 1), 'CTR', '00123', '00090', '客户子公司', '00045', 100, date(2026, 10, 1), '取消字样只是备注', '客户箱唛', '要求A'])
    s.append(['取消单'])
    s.append([None, 'CTR2', '00124', '00091', '子公司', '00046', 50, None])
    s.append(['已走货订单'])
    s.append([None, 'CTR3', '00125', '00092', '子公司', '00047', 80, None])
    b = BytesIO(); w.save(b)
    return b.getvalue()


def parsed(db, content=None):
    return history.preview(db, content or workbook(), '历史.xlsx', 'huaxing', 'buzzbee', 'BuzzBee')


def selections(p, amounts=('20', '10', '80')):
    return [HistorySelection(id=r['id'], status='cancelled' if r['section'] == 'cancelled' else 'active', opening_shipped_quantity=a)
            for r, a in zip(p['rows'], amounts)]


def confirm(db, p=None, **kwargs):
    p = p or parsed(db)
    args = dict(parsed=p, content=workbook(), selections=selections(p), fingerprint=p['fingerprint'], cutoff_date=date(2025, 12, 31), reason='期初逐单核对', actor='业务员')
    args.update(kwargs)
    return api.transaction(db, lambda: history.confirm(db, **args))


def test_restore_historical_cancellation_preserves_opening_and_original_fields(db):
    result = confirm(db)
    line = db.get(Line, result["items"][1]["id"])
    original = deepcopy(line.data)
    assert line.status == "cancelled" and line.shipped_quantity == 10
    api.transaction(db, lambda: ledger.restore(db, line,
        RestoreIn(expected_revision=line.revision, reason="原排期误取消恢复", confirmed=True), "业务员"))
    assert line.status == "active" and line.version == 2
    assert line.data == original
    assert ledger.line_out(db, line)["remaining_quantity"] == "40"
    assert db.scalar(select(func.count()).select_from(Shipment)) == 0


def test_parser_preserves_identifiers_custom_fields_and_section(db):
    p = parsed(db)
    assert [r['section'] for r in p['rows']] == ['unshipped', 'cancelled', 'shipped']
    r = p['rows'][0]
    assert not r['blocked'] and r['reference_no'] == '00123' and r['product_no'] == '00045'
    assert r['data']['carton_mark'] == '客户箱唛'
    assert r['data']['history_fields'][-1]['value'] == '要求A'
    assert r['data']['lineage']['quantity'].endswith('第 4 行')


def test_opening_is_audited_without_fake_shipments_or_auto_dispatch(db):
    p = parsed(db)
    result = confirm(db, p)
    assert result['created_count'] == 3
    assert [r['remaining_quantity'] for r in result['items']] == ['80', '0', '0']
    assert [r['shipped_quantity'] for r in result['items']] == ['20', '10', '80']
    assert db.scalar(select(func.count()).select_from(Shipment)) == 0
    assert db.scalar(select(func.count()).select_from(Dispatch)) == 0
    assert db.scalar(select(func.count()).select_from(Version)) == 3
    source = db.scalar(select(Source))
    assert source.content == workbook()  # unchanged original retained
    again = confirm(db, p)
    assert again['created_count'] == 0 and again['existing_count'] == 3


@pytest.mark.parametrize('amount', ['', '-1', 'NaN', 'Infinity', '100.00001', '101'])
def test_invalid_baselines_and_failed_batch_never_partially_persist(db, amount):
    p = parsed(db)
    with pytest.raises(Exception):
        confirm(db, p, selections=selections(p, ('20', amount, '80')))
    assert db.scalar(select(func.count()).select_from(Line)) == 0
    assert db.scalar(select(func.count()).select_from(Source)) == 0


def test_confirmation_rejects_changed_file_and_foreign_row(db):
    p = parsed(db)
    with pytest.raises(HTTPException):
        confirm(db, p, fingerprint='tampered')
    with pytest.raises(HTTPException):
        confirm(db, p, selections=[HistorySelection(id='other', status='active', opening_shipped_quantity='0')])
    assert db.scalar(select(func.count()).select_from(Line)) == 0


def test_migration_cannot_overwrite_existing_or_merge_duplicate_identity(db):
    p = parsed(db)
    confirm(db, p)
    with pytest.raises(HTTPException) as exc:
        confirm(db, p, selections=selections(p, ('0', '10', '80')))
    assert exc.value.status_code == 409
    assert parsed(db)['summary']['existing'] == 3
    w = openpyxl.load_workbook(BytesIO(workbook()))
    w.active.append([None, 'CTR', '00123', '00090', '子公司', '00045', 100])
    b = BytesIO(); w.save(b)
    dup = parsed(db, b.getvalue())
    assert dup['rows'][0]['blocked'] and dup['rows'][-1]['blocked']


def test_formula_cache_and_simulated_rows_block_without_guessing(db):
    w = openpyxl.load_workbook(BytesIO(workbook()))
    w.active['G4'] = '=50+50'
    w.active['I6'] = '模拟数据'
    b = BytesIO(); w.save(b)
    p = parsed(db, b.getvalue())
    assert p['rows'][0]['blocked'] and p['rows'][1]['blocked']
    assert any('公式' in x for x in p['rows'][0]['issues'])


def test_review_disagreement_is_reported_without_changing_item_status(db):
    w = openpyxl.load_workbook(BytesIO(workbook()))
    s = w.create_sheet('正单评审表')
    s.append([]); s.append([])
    s.append(['SO#/Reference', 'P/O#:', '产品编号', '數量'])
    s.append(['已走货订单'])
    s.append(['00123', '00090', '00045', 90])
    b = BytesIO(); w.save(b)
    p = parsed(db, b.getvalue())
    assert p['rows'][0]['blocked']
    assert p['rows'][0]['section'] == 'unshipped'
    assert any('数量不一致' in issue for issue in p['rows'][0]['issues'])


def test_opening_correction_preserves_later_shipments_and_original_version(db):
    result = confirm(db)
    line = db.get(Line, result['items'][0]['id'])
    ledger.ship(db, line, ship_body(line, qty='30'), '业务员')
    db.commit()
    original = db.scalar(select(Version).where(Version.line_id == line.id, Version.version == 1))
    history.correct_opening(db, line, HistoryOpeningIn(expected_revision=line.revision, opening_shipped_quantity='25', reason='核对原始排期'), '复核员')
    db.commit()
    assert line.shipped_quantity == Decimal('55')
    assert original.data['history_migration']['opening_shipped_quantity'] == '20'
    shipment = db.scalar(select(Shipment).where(Shipment.line_id == line.id))
    ledger.reverse_shipment(db, line, shipment, ReasonIn(expected_revision=line.revision, reason='出货记录更正'), '业务员')
    db.commit()
    assert line.shipped_quantity == Decimal('25')
    with pytest.raises(HTTPException):
        ledger.ship(db, line, ship_body(line, qty='1', key='earlier-01', document='OLD').model_copy(update={'ship_date': date(2025, 12, 31)}), '业务员')
    db.rollback()


def test_formal_po_reconciliation_keeps_historical_evidence(db):
    result = confirm(db)
    line = db.get(Line, result['items'][0]['id'])
    data = deepcopy(line.data)
    data.pop('history_migration'); data.pop('history_fields')
    data['reference_no'] = 'FORMAL-NEW'
    p = {'factory_id': 'huaxing', 'customer_code': 'buzzbee', 'customer_name': 'BuzzBee', 'rows': [data]}
    ledger.import_preview(db, preview=p, files=[('new.pdf', b'pdf', 'formal')], source_kind='formal', actor='a', reason='补齐正式订单', controls={}, reconcile_line=line, expected_revision=line.revision)
    db.commit()
    assert line.shipped_quantity == 20
    assert line.data['history_migration']['opening_shipped_quantity'] == '20'


def test_api_preview_is_read_only_confirmation_authorized_and_reparses(client, db):
    raw = workbook()
    form = {'factory_id': 'huaxing', 'customer_code': 'buzzbee'}
    files = {'schedule_file': ('历史.xlsx', raw, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
    response = client.post('/api/customer-order-ledger/history/preview', data=form, files=files)
    assert response.status_code == 200, response.text
    p = response.json()
    assert db.scalar(select(func.count()).select_from(Source)) == 0
    form.update(fingerprint=p['fingerprint'], cutoff_date='2025-12-31', reason='逐单核对期初', confirmed='true',
                selections=json.dumps([s.model_dump() for s in selections(p)]))
    result = client.post('/api/customer-order-ledger/history/confirm', data=form, files=files)
    assert result.status_code == 200, result.text
    assert result.json()['created_count'] == 3
    assert client.post('/api/customer-order-ledger/history/confirm', data=form, files=files).json()['existing_count'] == 3
    client.test_app.dependency_overrides[get_current_user] = lambda: user(['read', 'write'])
    assert client.post('/api/customer-order-ledger/history/confirm', data=form, files=files).status_code == 403
    assert client.get('/api/customer-order-ledger/history/customers?factory_id=huadeng').status_code == 403
    form['customer_code'] = 'jakks'
    assert client.post('/api/customer-order-ledger/history/preview', data=form, files=files).status_code == 400


def test_api_correction_requires_revision_and_permission(client, db):
    result = confirm(db)
    line = result['items'][0]
    url = f"/api/customer-order-ledger/lines/{line['id']}/history-opening?factory_id=huaxing"
    payload = {'expected_revision': 1, 'opening_shipped_quantity': '5', 'reason': '重新核对期初'}
    assert client.post(url, json=payload).status_code == 200
    assert client.post(url, json=payload).status_code == 409
    client.test_app.dependency_overrides[get_current_user] = lambda: user(['read', 'write'])
    assert client.post(url, json=payload).status_code == 403


def test_downstream_receives_balances_and_corrections_without_private_history(db):
    result = confirm(db)
    line = db.get(Line, result['items'][0]['id'])
    ledger.dispatch(db, line, DispatchIn(expected_revision=1, recipients=['pmc']), '业务员')
    db.commit()
    first = db.scalar(select(Dispatch).where(Dispatch.line_id == line.id, Dispatch.version == 1))
    assert first.snapshot['quantity'] == '100'
    assert first.snapshot['shipped_quantity'] == '20'
    assert first.snapshot['remaining_quantity'] == '80'
    assert first.snapshot['history_cutoff_date'] == '2025-12-31'
    assert 'history_fields' not in first.snapshot and 'history_migration' not in first.snapshot
    history.correct_opening(db, line, HistoryOpeningIn(expected_revision=line.revision, opening_shipped_quantity='30', reason='核对实际期初'), '复核员')
    db.commit()
    second = db.scalar(select(Dispatch).where(Dispatch.line_id == line.id, Dispatch.version == 2))
    assert second.snapshot['remaining_quantity'] == '70'
    assert first.snapshot['remaining_quantity'] == '80'
