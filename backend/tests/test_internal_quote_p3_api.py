import hashlib
import importlib
import json
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as WorksheetImage
from openpyxl.styles import Font
from PIL import Image as PillowImage
from docx import Document

from app.services.internal_quote_excel import ENGINEERING_WORKBOOK_TEMPLATE_PATH
from test_internal_quote_api import ALL_SECTION_CODES, create_payload, login, logout, make_client


def workbook_bytes(rows: list[list[object]], title: str = "报价明细") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def document_bytes(paragraphs: list[str]) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def workbook_bytes_with_image(rows: list[list[object]], anchor: str, title: str = "报价明细") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    image_bytes = BytesIO()
    PillowImage.new("RGB", (4, 4), color=(16, 118, 110)).save(image_bytes, format="PNG")
    image_bytes.seek(0)
    image = WorksheetImage(image_bytes)
    image.anchor = anchor
    sheet.add_image(image)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def workbook_bytes_with_linked_sheets() -> bytes:
    workbook = Workbook()
    detail = workbook.active
    detail.title = "报价明细"
    detail["A1"] = "上传源表"
    detail["A1"].font = Font(bold=True, color="FF0000")
    detail.merge_cells("A1:B1")
    parameter = workbook.create_sheet("参数")
    parameter["A1"] = "='报价明细'!A1"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_p3_import_preview_is_non_mutating_and_confirm_is_revision_locked(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "iq_p3_creator", "sales_customer_owner", "sales-business")
        created = client.post(
            "/api/internal-quotes",
            json=create_payload(suffix="P3-IMPORT", participating_sections=ALL_SECTION_CODES),
        ).json()
        quote_id = created["id"]
        source = workbook_bytes_with_image(
            [
                ["模号", "产品名称", "材质", "克重", "套数", "模价", "机型", "目标数"],
                ["M-300", "主体模", "ABS", 100, 1, 7750, "4A", 5000],
            ],
            "U2",
        )

        sales_preview = client.post(
            f"/api/internal-quotes/{quote_id}/imports/mold/preview",
            files={"file": ("模具报价.xlsx", source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert sales_preview.status_code == 201, sales_preview.text
        assert sales_preview.json()["target_department"] == "engineering"

        logout(client)
        login(client, "iq_p3_engineer", "engineer", "engineering")
        preview_response = client.post(
            f"/api/internal-quotes/{quote_id}/imports/mold/preview",
            files={"file": ("模具报价.xlsx", source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert preview_response.status_code == 201, preview_response.text
        preview = preview_response.json()
        assert preview["status"] == "previewed"
        assert preview["target_department"] == "engineering"
        assert preview["target_revision"] == 1
        assert preview["row_count"] == 1
        assert preview["diff_summary"] == {
            "target_revision": 1,
            "existing_rows": 0,
            "imported_rows": 1,
            "replace_result_rows": 1,
            "append_result_rows": 1,
        }
        assert preview["source_sha256"] == hashlib.sha256(source).hexdigest()
        assert preview["payload_fragment"]["molds"][0]["image_reference"] == "模具图片-U2-1.png"
        assert "embedded_images" not in preview

        detail_before = client.get(f"/api/internal-quotes/{quote_id}").json()
        engineering_before = next(
            item for item in detail_before["sections"] if item["department"] == "engineering"
        )
        assert engineering_before["revision"] == 1
        assert engineering_before["payload"] == {}

        confirm_response = client.post(
            f"/api/internal-quotes/{quote_id}/imports/{preview['batch_id']}/confirm",
            json={"revision": 1, "mode": "replace"},
        )
        assert confirm_response.status_code == 200, confirm_response.text
        confirmed = confirm_response.json()
        assert confirmed["batch"]["status"] == "confirmed"
        assert confirmed["batch"]["confirm_mode"] == "replace"
        assert confirmed["section"]["revision"] == 2
        assert confirmed["section"]["calculation_status"] == "valid"
        assert confirmed["section"]["payload"]["molds"][0]["mold_no"] == "M-300"
        attachment_ids = confirmed["section"]["payload"]["molds"][0]["image_attachment_ids"]
        assert len(attachment_ids) == 1
        assert confirmed["section"]["calculation"]["totals"]["mold_total_rmb"] == "7750.0000"

        attachments = client.get(
            f"/api/internal-quotes/{quote_id}/attachments?department=engineering"
        ).json()
        assert {item["id"] for item in attachments} == {
            *attachment_ids,
            next(item["id"] for item in attachments if item["file_name"] == "模具报价.xlsx"),
        }
        assert {item["file_name"] for item in attachments} == {
            "模具报价.xlsx",
            "模具图片-U2-1.png",
        }
        source_attachment = next(
            item for item in attachments if item["file_name"] == "模具报价.xlsx"
        )
        workbook_preview = client.get(
            f"/api/internal-quotes/{quote_id}/attachments/{source_attachment['id']}/preview"
        )
        assert workbook_preview.status_code == 200
        assert workbook_preview.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert workbook_preview.headers["content-disposition"].startswith("inline;")
        assert workbook_preview.content == source
        workbook_content_preview = client.get(
            f"/api/internal-quotes/{quote_id}/attachments/{source_attachment['id']}/content-preview"
        )
        assert workbook_content_preview.status_code == 200
        assert workbook_content_preview.json()["kind"] == "excel"
        assert workbook_content_preview.json()["sheets"][0]["name"] == "报价明细"
        assert workbook_content_preview.json()["sheets"][0]["rows"][0][:2] == ["模号", "产品名称"]

        word_source = document_bytes(["工程资料说明", "只用于当前部门核价。"])
        word_upload = client.post(
            f"/api/internal-quotes/{quote_id}/attachments",
            data={"department": "engineering"},
            files={"file": ("工程说明.docx", word_source, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert word_upload.status_code == 201, word_upload.text
        word_content_preview = client.get(
            f"/api/internal-quotes/{quote_id}/attachments/{word_upload.json()['id']}/content-preview"
        )
        assert word_content_preview.status_code == 200
        assert word_content_preview.json()["kind"] == "word"
        assert word_content_preview.json()["paragraphs"] == ["工程资料说明", "只用于当前部门核价。"]
        image_preview = client.get(
            f"/api/internal-quotes/{quote_id}/attachments/{attachment_ids[0]}/preview"
        )
        assert image_preview.status_code == 200
        assert image_preview.headers["content-type"] == "image/png"
        assert image_preview.headers["content-disposition"].startswith("inline;")
        assert image_preview.content.startswith(b"\x89PNG\r\n\x1a\n")

        duplicate = client.post(
            f"/api/internal-quotes/{quote_id}/imports/{preview['batch_id']}/confirm",
            json={"revision": 2, "mode": "append"},
        )
        assert duplicate.status_code == 409

        batches = client.get(f"/api/internal-quotes/{quote_id}/imports")
        assert batches.status_code == 200
        listed_batch = next(
            item for item in batches.json()
            if item["batch_id"] == preview["batch_id"]
        )
        assert listed_batch["confirmed_revision"] == 2


def test_hardware_template_preview_maps_only_shared_material_fields_and_replace_keeps_auxiliary(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "iq_p3_hardware_creator", "sales_customer_owner", "sales-business")
        created = client.post(
            "/api/internal-quotes",
            json=create_payload(suffix="P3-HARDWARE", participating_sections=ALL_SECTION_CODES),
        ).json()
        quote_id = created["id"]

        logout(client)
        login(client, "iq_p3_hardware_engineer", "engineer", "engineering")
        saved = client.put(
            f"/api/internal-quotes/{quote_id}/sections/engineering",
            json={
                "revision": 1,
                "payload": {
                    "materials": [
                        {"item": "旧五金", "category": "hardware", "quantity": 1, "unit_price_rmb": 1},
                        {"item": "胶水", "category": "auxiliary", "quantity": 2, "unit_price_rmb": 3},
                    ],
                    "molds": [],
                    "amortization_qty": 0,
                    "customer_mold_subsidy_usd": 0,
                },
            },
        )
        assert saved.status_code == 200, saved.text

        source = workbook_bytes(
            [
                ["(五金电子部分)"],
                ["货号", "HW-01"],
                ["产品名称", "测试产品"],
                [],
                [],
                ["零件名称", "配件用处", "规格", "用量", "单位", "单价RMB", "总价", "质料", "表面处理", "供应商", "联系人/电话", "备注", "图片"],
                [],
                ["螺丝", "尿兜", "2.6*8PB", 3, "pcs", 0.0044, 0.0132, "铁", "镀镍", "港正", "张/13800000000", "样板", ""],
                ["合计", "", "", "", "", "", 0.0132],
                ["附：五金尺寸标准表（其它配件及包装物料请参照客签板）"],
            ],
            title="外购",
        )
        preview_response = client.post(
            f"/api/internal-quotes/{quote_id}/imports/hardware/preview",
            files={"file": ("五金1.xlsx", source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert preview_response.status_code == 201, preview_response.text
        preview = preview_response.json()
        assert preview["header_row"] == 6
        assert preview["row_count"] == 1
        assert preview["diff_summary"]["existing_rows"] == 1
        imported = preview["payload_fragment"]["materials"][0]
        assert imported == {
            "item": "螺丝",
            "category": "hardware",
            "specification": "2.6*8PB",
            "quantity": "3.0000",
            "unit_price_rmb": "0.0044",
            "auxiliary_category": "五金",
            "tax_rate_percent": "13.0000",
            "remark": "样板",
            "source_row": 8,
        }
        assert not {"purpose", "unit", "material", "surface_treatment", "supplier", "contact"}.intersection(imported)
        assert any("模板其他列不写入报价字段" in warning for warning in preview["warnings"])

        confirmed = client.post(
            f"/api/internal-quotes/{quote_id}/imports/{preview['batch_id']}/confirm",
            # Older clients may still submit append. The server owns the new
            # deterministic contract and must replace the template region.
            json={"revision": 2, "mode": "append"},
        )
        assert confirmed.status_code == 200, confirmed.text
        assert confirmed.json()["batch"]["confirm_mode"] == "replace"
        section = confirmed.json()["section"]
        materials = section["payload"]["materials"]
        assert [row["item"] for row in materials] == ["胶水", "螺丝"]
        assert section["calculation"]["totals"]["hardware_hkd"] == "0.0155"
        hardware_line = next(row for row in section["calculation"]["line_breakdown"] if row["category"] == "hardware")
        assert hardware_line["unit_price_hkd"] == "0.0052"
        assert hardware_line["auxiliary_category"] == "五金"
        assert hardware_line["tax_rate_percent"] == "13.0000"
        assert "冻结 RMB→HKD 汇率" in hardware_line["formula"]


def test_import_source_attachment_delete_clears_generated_data_and_allows_clean_reimport(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "iq_p3_delete_import_creator", "sales_customer_owner", "sales-business")
        created = client.post(
            "/api/internal-quotes",
            json=create_payload(suffix="P3-DELETE-IMPORT", participating_sections=ALL_SECTION_CODES),
        ).json()
        quote_id = created["id"]

        logout(client)
        login(client, "iq_p3_delete_import_engineer", "engineer", "engineering")
        source = workbook_bytes_with_image(
            [
                ["模号", "产品名称", "材质", "克重", "套数", "模价", "机型", "目标数"],
                ["M-DELETE", "待删除主体模", "ABS", 100, 1, 7750, "4A", 5000],
            ],
            "U2",
        )
        preview = client.post(
            f"/api/internal-quotes/{quote_id}/imports/mold/preview",
            files={"file": ("待删除模具报价.xlsx", source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).json()
        confirmed = client.post(
            f"/api/internal-quotes/{quote_id}/imports/{preview['batch_id']}/confirm",
            json={"revision": 1},
        )
        assert confirmed.status_code == 200, confirmed.text
        assert confirmed.json()["batch"]["confirm_mode"] == "replace"
        imported_row = confirmed.json()["section"]["payload"]["molds"][0]
        assert imported_row["import_batch_id"] == preview["batch_id"]
        image_attachment_ids = imported_row["image_attachment_ids"]

        attachments = client.get(
            f"/api/internal-quotes/{quote_id}/attachments?department=engineering"
        ).json()
        source_attachment = next(item for item in attachments if item["is_import_source"])
        assert source_attachment["file_name"] == "待删除模具报价.xlsx"
        assert source_attachment["import_batch_id"] == preview["batch_id"]
        assert source_attachment["import_type"] == "mold"

        stale_delete = client.delete(
            f"/api/internal-quotes/{quote_id}/attachments/{source_attachment['id']}",
            params={"revision": 1},
        )
        assert stale_delete.status_code == 409

        deleted = client.delete(
            f"/api/internal-quotes/{quote_id}/attachments/{source_attachment['id']}",
            params={"revision": 2},
        )
        assert deleted.status_code == 204, deleted.text

        detail = client.get(f"/api/internal-quotes/{quote_id}").json()
        engineering = next(item for item in detail["sections"] if item["department"] == "engineering")
        assert engineering["revision"] == 3
        assert engineering["status"] == "draft"
        assert engineering["payload"]["molds"] == []
        remaining_attachments = client.get(
            f"/api/internal-quotes/{quote_id}/attachments?department=engineering"
        ).json()
        assert source_attachment["id"] not in {item["id"] for item in remaining_attachments}
        assert not set(image_attachment_ids).intersection(item["id"] for item in remaining_attachments)
        batches = client.get(f"/api/internal-quotes/{quote_id}/imports").json()
        assert next(item for item in batches if item["batch_id"] == preview["batch_id"])["status"] == "deleted"
        timeline = client.get(f"/api/internal-quotes/{quote_id}/timeline").json()
        assert any(item["action"] == "delete_import_attachment" for item in timeline["business_events"])

        reimport_preview = client.post(
            f"/api/internal-quotes/{quote_id}/imports/mold/preview",
            files={"file": ("待删除模具报价.xlsx", source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        ).json()
        reimported = client.post(
            f"/api/internal-quotes/{quote_id}/imports/{reimport_preview['batch_id']}/confirm",
            json={"revision": 3},
        )
        assert reimported.status_code == 200, reimported.text
        assert reimported.json()["section"]["payload"]["molds"][0]["mold_no"] == "M-DELETE"


def test_molding_quote_preview_and_confirm_recalculate_from_frozen_reference(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "iq_p3_molding_creator", "sales_customer_owner", "sales-business")
        created = client.post(
            "/api/internal-quotes",
            json=create_payload(suffix="P3-MOLDING", participating_sections=ALL_SECTION_CODES),
        ).json()
        quote_id = created["id"]

        logout(client)
        login(client, "iq_p3_molding_clerk", "molding_clerk", "molding")
        source = workbook_bytes(
            [
                ["模具名称", "模号", "材质", "料型", "颜色", "啤净重(g)", "料损耗 3%", "料价 HK$/g", "原料单价 HK$", "机台", "啤价(HK$/啤)", "出模数", "套数", "机型", "目标数", "周期(秒)", "成品金额 HK$"],
                ["主体模", "M-01", "ABS", "750SW", "黑色", 100, 103, 999, 999, "80T", 999, "2", 1, "5A", 5000, 24, 999],
                [],
                ["货名", "日产量/22H", "用料", "预估料重 g", "料价 HK$/lb", "产品料价", "吹工", "披锋", "小计", "利润 ×", "合计 HK$", "出数", "模价 (¥)"],
                ["吹气瓶", "12000", "ABS 750SW", 45, 999, 999, 0.2, 0.1, 999, 1.05, 999, "1出2", 5000],
            ],
            title="啤机报价",
        )
        preview_response = client.post(
            f"/api/internal-quotes/{quote_id}/imports/molding/preview",
            files={"file": ("啤机报价.xlsx", source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert preview_response.status_code == 201, preview_response.text
        preview = preview_response.json()
        assert preview["target_department"] == "molding"
        assert preview["row_count"] == 2
        assert preview["diff_summary"]["existing_rows"] == 0
        assert any("不直接写入" in warning for warning in preview["warnings"])

        confirmed = client.post(
            f"/api/internal-quotes/{quote_id}/imports/{preview['batch_id']}/confirm",
            json={"revision": 1, "mode": "replace"},
        )
        assert confirmed.status_code == 200, confirmed.text
        section = confirmed.json()["section"]
        assert section["revision"] == 2
        assert section["calculation_status"] == "valid"
        assert section["payload"]["injection_lines"][0]["mold_no"] == "M-01"
        assert section["payload"]["blow_lines"][0]["daily_capacity"] == "12000"
        injection = section["calculation"]["line_breakdown"][0]
        assert injection["material_price_hkd_lb"] == "8.5000"
        assert injection["machine_shift_price_hkd"] == "940.0000"
        assert injection["unit_amount_hkd"] != "999.0000"


def test_assembly_workshop_preview_and_confirm_maps_regions_and_recalculates(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "iq_p3_assembly_creator", "sales_customer_owner", "sales-business")
        created = client.post(
            "/api/internal-quotes",
            json=create_payload(suffix="P3-ASSEMBLY", participating_sections=ALL_SECTION_CODES),
        ).json()
        quote_id = created["id"]

        logout(client)
        login(client, "iq_p3_assembly_admin", "admin", "assembly")
        source = workbook_bytes(
            [
                ["报客价/港币", "报价人", "货号或图片", "做工名称", "总目标数量", "人数"],
                ["车间填写", "车间填写", "车间填写", "车间填写", "车间填写", "车间填写"],
                [None, 8, "组装桶", "测试IC板", 3000, 1],
                [None, None, None, "焊喇叭", 3000, 2],
                [None, None, None, None, None, 3],
                [None, 9, "包装公仔", "彩盒印日期码", 3000, 4],
                [None, None, None, None, None, 4],
            ],
            title="组装",
        )
        preview_response = client.post(
            f"/api/internal-quotes/{quote_id}/imports/assembly/preview",
            files={
                "file": (
                    "装工.xlsx",
                    source,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert preview_response.status_code == 201, preview_response.text
        preview = preview_response.json()
        assert preview["target_department"] == "assembly"
        assert preview["header_row"] == 1
        assert preview["row_count"] == 3
        assert [
            (group["name"], group["category"], group["production_qty"])
            for group in preview["payload_fragment"]["groups"]
        ] == [
            ("组装桶", "assembly", "3000.0000"),
            ("包装公仔", "packaging", "3000.0000"),
        ]

        confirmed = client.post(
            f"/api/internal-quotes/{quote_id}/imports/{preview['batch_id']}/confirm",
            json={"revision": 1, "mode": "replace"},
        )
        assert confirmed.status_code == 200, confirmed.text
        section = confirmed.json()["section"]
        assert section["revision"] == 2
        assert section["calculation_status"] == "valid"
        assert section["calculation"]["totals"] == {
            "assembly_hkd": "0.2600",
            "packaging_hkd": "0.3467",
            "total_hkd": "0.6067",
        }
        assert section["calculation"]["group_summaries"][0]["total_persons"] == "3.0000"
        assert section["calculation"]["group_summaries"][1]["total_persons"] == "4.0000"


def test_p3_attachment_validates_magic_deduplicates_and_downloads(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "iq_p3_attachment_creator", "sales_customer_owner", "sales-business")
        quote = client.post(
            "/api/internal-quotes",
            json=create_payload(suffix="P3-ATTACHMENT", participating_sections=ALL_SECTION_CODES),
        ).json()
        quote_id = quote["id"]

        logout(client)
        login(client, "iq_p3_attachment_engineer", "engineer", "engineering")
        image = b"\x89PNG\r\n\x1a\n" + b"internal-quote-image"
        uploaded = client.post(
            f"/api/internal-quotes/{quote_id}/attachments",
            data={"department": "engineering"},
            files={"file": ("模具图.png", image, "image/png")},
        )
        assert uploaded.status_code == 201, uploaded.text
        attachment = uploaded.json()
        assert attachment["sha256"] == hashlib.sha256(image).hexdigest()
        assert attachment["content_type"] == "image/png"

        duplicate = client.post(
            f"/api/internal-quotes/{quote_id}/attachments",
            data={"department": "engineering"},
            files={"file": ("另一名称.png", image, "image/png")},
        )
        assert duplicate.status_code == 409

        invalid = client.post(
            f"/api/internal-quotes/{quote_id}/attachments",
            data={"department": "engineering"},
            files={"file": ("伪造图片.png", b"not-a-png", "image/png")},
        )
        assert invalid.status_code == 400

        listed = client.get(
            f"/api/internal-quotes/{quote_id}/attachments?department=engineering"
        )
        assert listed.status_code == 200
        assert [item["id"] for item in listed.json()] == [attachment["id"]]

        downloaded = client.get(
            f"/api/internal-quotes/{quote_id}/attachments/{attachment['id']}/download"
        )
        assert downloaded.status_code == 200
        assert downloaded.content == image
        assert downloaded.headers["x-content-sha256"] == attachment["sha256"]

        supporting_attachment_delete = client.delete(
            f"/api/internal-quotes/{quote_id}/attachments/{attachment['id']}",
            params={"revision": 1},
        )
        assert supporting_attachment_delete.status_code == 409
        assert "不是结构化报价导入源文件" in supporting_attachment_delete.json()["detail"]

        saved = client.put(
            f"/api/internal-quotes/{quote_id}/sections/engineering",
            json={
                "revision": 1,
                "payload": {
                    "materials": [
                        {"item": "五金件", "category": "hardware", "quantity": "2", "unit_price_rmb": "8.5"}
                    ],
                    "molds": [],
                    "production_mold_fees": [],
                },
            },
        )
        assert saved.status_code == 200, saved.text
        submitted = client.post(
            f"/api/internal-quotes/{quote_id}/sections/engineering/submit",
            json={"revision": 2},
        )
        assert submitted.status_code == 200, submitted.text
        late_upload = client.post(
            f"/api/internal-quotes/{quote_id}/attachments",
            data={"department": "engineering"},
            files={"file": ("提交后新增.png", image + b"late", "image/png")},
        )
        assert late_upload.status_code == 409
        assert "重开" in late_upload.json()["detail"]


def test_p3_production_attachment_reads_are_scoped_to_its_department(monkeypatch):
    with make_client(monkeypatch) as client:
        login(client, "iq_attachment_scope_owner", "sales_customer_owner", "sales-business")
        quote = client.post(
            "/api/internal-quotes",
            json=create_payload(suffix="ATTACHMENT-SCOPE", participating_sections=ALL_SECTION_CODES),
        ).json()
        quote_id = quote["id"]
        engineering_image = b"\x89PNG\r\n\x1a\n" + b"engineering-document"
        molding_image = b"\x89PNG\r\n\x1a\n" + b"molding-document"
        engineering_upload = client.post(
            f"/api/internal-quotes/{quote_id}/attachments",
            data={"department": "engineering"},
            files={"file": ("工程资料.png", engineering_image, "image/png")},
        )
        molding_upload = client.post(
            f"/api/internal-quotes/{quote_id}/attachments",
            data={"department": "molding"},
            files={"file": ("啤机资料.png", molding_image, "image/png")},
        )
        assert engineering_upload.status_code == 201, engineering_upload.text
        assert molding_upload.status_code == 201, molding_upload.text

        logout(client)
        login(client, "iq_attachment_scope_molding", "molding_clerk", "molding")
        listed = client.get(f"/api/internal-quotes/{quote_id}/attachments")
        assert listed.status_code == 200, listed.text
        assert [item["file_name"] for item in listed.json()] == ["啤机资料.png"]

        forbidden_list = client.get(
            f"/api/internal-quotes/{quote_id}/attachments?department=engineering"
        )
        assert forbidden_list.status_code == 403
        forbidden_preview = client.get(
            f"/api/internal-quotes/{quote_id}/attachments/"
            f"{engineering_upload.json()['id']}/preview"
        )
        assert forbidden_preview.status_code == 403
        own_preview = client.get(
            f"/api/internal-quotes/{quote_id}/attachments/"
            f"{molding_upload.json()['id']}/preview"
        )
        assert own_preview.status_code == 200
        assert own_preview.content == molding_image


def test_p3_controlled_export_is_retained_reproducible_and_superseded(monkeypatch):
    with make_client(monkeypatch) as client:
        profile = login(
            client,
            "iq_p3_exporter",
            "sales_customer_supervisor",
            "sales-business",
        )
        assert "internal_quote:export" in profile["permissions"]
        quote = client.post(
            "/api/internal-quotes",
            json=create_payload(suffix="P3-EXPORT", participating_sections=ALL_SECTION_CODES),
        ).json()
        quote_id = quote["id"]

        blocked = client.post(f"/api/internal-quotes/{quote_id}/exports")
        assert blocked.status_code == 409
        assert set(blocked.json()["detail"]["incomplete_sections"]) == {
            "sales",
            "engineering",
            "electronic",
            "molding",
            "painting",
            "slush",
            "sewing",
            "hair",
            "assembly",
        }
        source_attachment = workbook_bytes_with_linked_sheets()
        uploaded_source = client.post(
            f"/api/internal-quotes/{quote_id}/attachments",
            data={"department": "engineering"},
            files={
                "file": (
                    "工程核价依据.xlsx",
                    source_attachment,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert uploaded_source.status_code == 201, uploaded_source.text

        db_module = importlib.import_module("app.db")
        quote_models = importlib.import_module("app.models.internal_quote")
        with db_module.SessionLocal() as db:
            sections = db.query(quote_models.InternalQuoteSection).filter_by(quote_id=quote_id).all()
            for section in sections:
                section.status = "not_applicable"
                section.calculation_status = "not_applicable"
                section.dependency_status = "current"
                if section.department == "sales":
                    section.payload_json = json.dumps({
                        "paper_price_factor": 2.75,
                        "shipping": {
                            "markup_x": 1.99,
                            "markup_tiers": [
                                {"moq": 3000, "markup_x": 1.25},
                                {"moq": 5000, "markup_x": 1.15},
                                {"moq": 10000, "markup_x": 1.10},
                            ],
                            "selected_markup_moq": 10000,
                            "misc_ratio": 0.035,
                        },
                        "testing_fee_total_usd": 1500,
                        "testing_fee_moqs": [3000, 5000, 10000],
                        "packaging_materials": [{
                            "item": "彩盒",
                            "specification": "四彩印刷",
                            "quantity": 1,
                        }],
                        "product_size_in": {"length": 12, "width": 8, "height": 4},
                        "color_box_size_unit": "cm",
                        "color_box_size_in": {"length": 10, "width": 5, "height": 4},
                        "cartons": [{
                            "item": "主纸箱",
                            "size_unit": "cm",
                            "length_in": 20,
                            "width_in": 10,
                            "height_in": 8,
                            "qty_per_carton": 2,
                            "flat_cards": [],
                        }],
                    }, ensure_ascii=False)
                    section.calculation_json = json.dumps({
                        "status": "valid",
                        "line_breakdown": [],
                            "totals": {
                                "carton_hkd": "0.38",
                                "testing_fee_total_usd": "1500",
                            "testing_fee_tiers": [
                                {"moq": "3000", "unit_price_usd": "0.5"},
                                {"moq": "5000", "unit_price_usd": "0.3"},
                                {"moq": "10000", "unit_price_usd": "0.15"},
                            ],
                            "freight_options": [
                                {
                                    "route_key": "hk40",
                                    "item": "HK 40 柜",
                                    "has_lifting_fee": True,
                                    "freight_per_piece_hkd": "2.06",
                                    "lifting_per_piece_hkd": "3.71",
                                    "total_cartons": "1200",
                                },
                                {
                                    "route_key": "yt20",
                                    "item": "YT 20 柜",
                                    "has_lifting_fee": True,
                                    "freight_per_piece_hkd": "3.01",
                                    "lifting_per_piece_hkd": "1.97",
                                    "total_cartons": "600",
                                },
                            ],
                        },
                    }, ensure_ascii=False)
                    section.calculation_status = "valid"
                if section.department == "engineering":
                    section.payload_json = json.dumps({
                        "materials": [
                            {
                                "item": "螺丝",
                                "category": "hardware",
                                "specification": "2.6×8PB",
                                "material": "铁",
                                "quantity": 2,
                                "supplier": "港正",
                                "surface_treatment": "镀镍",
                            },
                            {
                                "item": "润滑油",
                                "category": "auxiliary",
                                "specification": "工程用",
                                "quantity": 0.02,
                            },
                        ],
                        "molds": [
                            {
                                "item": "左右前枪身（橙色）",
                                "mold_no": "M01",
                                "color": "橙色",
                                "cavity": "2",
                                "quantity": 1,
                            },
                            {
                                "item": "泵杆/击锤/扣机/配件(7件)",
                                "mold_no": "M05",
                                "color": "黑色",
                                "cavity": "4",
                                "quantity": 1,
                            },
                        ],
                    }, ensure_ascii=False)
                if section.department == "electronic":
                    section.payload_json = json.dumps({
                        "quote_mode": "detail",
                        "components": [{
                            "item": "主控板",
                            "specification": "PCB-A1",
                            "quantity": 1,
                            "children": [{
                                "item": "喇叭",
                                "specification": "8Ω",
                                "quantity": 1,
                            }],
                        }],
                    }, ensure_ascii=False)
                if section.department == "molding":
                    section.payload_json = json.dumps({
                        "injection_lines": [
                            {
                                "item": "左右前枪身（橙色）",
                                "mold_no": "M01",
                                "material": "ABS",
                                "grade": "KF740",
                                "color": "橙色",
                                "net_weight_g": 52.3,
                                "cavity": "2",
                                "sets": 1,
                                "machine_code": "20A",
                                "target_output": 5000,
                            },
                            {
                                "item": "泵杆/击锤/扣机/配件(7件)",
                                "mold_no": "M05",
                                "material": "ABS",
                                "grade": "KF740",
                                "color": "黑色",
                                "net_weight_g": 44.8,
                                "cavity": "4",
                                "sets": 1,
                                "machine_code": "18A",
                                "target_output": 4200,
                            },
                        ],
                    }, ensure_ascii=False)
                if section.department == "assembly":
                    section.status = "approved"
                    section.calculation_status = "valid"
                    section.payload_json = json.dumps(
                        {
                            "groups": [
                                {"name": "组装马桶", "category": "assembly"},
                                {"name": "奶瓶", "category": "assembly"},
                                {"name": "勺子", "category": "assembly"},
                                {"name": "瓶子", "category": "assembly"},
                                {"name": "组装公仔", "category": "assembly"},
                                {"name": "包装公仔", "category": "packaging"},
                            ]
                        },
                        ensure_ascii=False,
                    )
                    section.calculation_json = json.dumps(
                        {
                            "status": "valid",
                            "line_breakdown": [],
                            "group_summaries": [
                                {
                                    "category": "assembly",
                                    "group": "组装马桶",
                                    "standard_work_hours": "11.0000",
                                    "production_qty": "3000.0000",
                                    "total_persons": "17.0000",
                                    "amount_hkd_pcs": "1.4733",
                                },
                                {
                                    "category": "assembly",
                                    "group": "奶瓶",
                                    "standard_work_hours": "11.0000",
                                    "production_qty": "3000.0000",
                                    "total_persons": "7.0000",
                                    "amount_hkd_pcs": "0.6067",
                                },
                                {
                                    "category": "assembly",
                                    "group": "勺子",
                                    "standard_work_hours": "11.0000",
                                    "production_qty": "6000.0000",
                                    "total_persons": "5.0000",
                                    "amount_hkd_pcs": "0.2167",
                                },
                                {
                                    "category": "assembly",
                                    "group": "瓶子",
                                    "standard_work_hours": "11.0000",
                                    "production_qty": "3000.0000",
                                    "total_persons": "13.0000",
                                    "amount_hkd_pcs": "1.1267",
                                },
                                {
                                    "category": "assembly",
                                    "group": "组装公仔",
                                    "standard_work_hours": "11.0000",
                                    "production_qty": "3000.0000",
                                    "total_persons": "13.0000",
                                    "amount_hkd_pcs": "1.1267",
                                },
                                {
                                    "category": "packaging",
                                    "group": "包装公仔",
                                    "standard_work_hours": "11.0000",
                                    "production_qty": "3000.0000",
                                    "total_persons": "40.0000",
                                    "amount_hkd_pcs": "3.4667",
                                },
                            ],
                            "totals": {
                                "assembly_hkd": "4.5501",
                                "packaging_hkd": "3.4667",
                                "total_hkd": "8.0168",
                            },
                        },
                        ensure_ascii=False,
                    )
                db.commit()

        first_response = client.post(f"/api/internal-quotes/{quote_id}/exports")
        assert first_response.status_code == 201, first_response.text
        first = first_response.json()
        assert first["status"] == "current"
        assert first["template_version"] == "internal-quote-p3-v1"
        assert first["release_stage"] == "p3_section_approved"
        assert first["export_manifest"]["p4_final_release_required"] is True
        assert first["export_manifest"]["workbook_layout_version"] == "internal-quote-unified-desk-v28"
        assert first["export_manifest"]["export_file_name_version"] == "quote-product-date-v1"
        assert first["file_name"] == (
            f"{quote['quote_no']}_{quote['product_name']}_{first['exported_at'][:10]}.xlsx"
        )
        assert first["export_manifest"]["spreadsheet_attachments"][0]["file_name"] == "工程核价依据.xlsx"

        download = client.get(
            f"/api/internal-quotes/{quote_id}/exports/{first['id']}/download"
        )
        assert download.status_code == 200
        assert download.content.startswith(b"PK")
        assert hashlib.sha256(download.content).hexdigest() == first["sha256"]
        workbook = load_workbook(BytesIO(download.content), data_only=False, read_only=True)
        assert workbook.sheetnames == [
            "报价明细",
            "电子明细",
            "车缝明细",
            "车发明细",
            "装配明细",
            "审批与版本",
            "工程核价依据-报价明细",
            "工程核价依据-参数",
        ]
        quote_sheet = workbook["报价明细"]
        assert quote_sheet.sheet_state == "visible"
        assert quote_sheet["A7"].value == "内部报价测试产品报价"
        packaging_labels = [quote_sheet.cell(row, 14).value for row in range(1, quote_sheet.max_row + 1)]
        assert "彩盒尺寸 (cm)" in packaging_labels
        assert "产品尺寸 (in)" in packaging_labels
        carton_row = next(row for row in range(1, quote_sheet.max_row + 1) if quote_sheet.cell(row, 14).value == "外箱 (cm):")
        assert quote_sheet.cell(carton_row, 15).value == 50.8
        assert quote_sheet.cell(carton_row + 1, 15).value == 25.4
        assert quote_sheet.cell(carton_row + 3, 15).value == f"=O{carton_row}/2.54*P{carton_row}/2.54*Q{carton_row}/2.54/1728"
        assert [quote_sheet.cell(8, column).value for column in range(3, 13)] == [
            "名称", "料型", "料重(G)", "料价(G)", "机型", "1出几套", "目标数", "啤工", "料金额", "报价啤工",
        ]
        route_header_row = next(
            row
            for row in range(1, quote_sheet.max_row + 1)
            if quote_sheet.cell(row, 2).value == "运输方案"
        )
        assert [quote_sheet.cell(route_header_row, column).value for column in (5, 6)] == ["HK 40 柜", "YT 20 柜"]
        assert [quote_sheet.cell(route_header_row + 1, column).value for column in (5, 6)] == [2.06, 3.01]
        assert [quote_sheet.cell(route_header_row + 2, column).value for column in (5, 6)] == [3.71, 1.97]
        subtotal_row = route_header_row + 3
        assert quote_sheet.cell(subtotal_row, 4).value == f"=SUM(D19:D{route_header_row})"
        quote_rows = {
            quote_sheet.cell(row, 2).value: row
            for row in range(1, quote_sheet.max_row + 1)
            if str(quote_sheet.cell(row, 2).value or "").startswith("报价（MOQ")
        }
        assert set(quote_rows) == {"报价（MOQ3K）", "报价（MOQ5K）", "报价（MOQ10K）"}
        for quote_row in quote_rows.values():
            settlement_row = quote_row - 1
            assert quote_sheet.cell(settlement_row, 3).value == "÷"
            assert all(
                quote_sheet.cell(settlement_row, column).value == "=1-$Q$6"
                and quote_sheet.cell(settlement_row, column).number_format == "0.0000"
                for column in (4, 5, 6)
            )
        assert quote_sheet.cell(quote_rows["报价（MOQ3K）"], 4).value == (
            f"=D{subtotal_row}*D{quote_rows['报价（MOQ3K）'] - 2}/D{quote_rows['报价（MOQ3K）'] - 1}"
        )
        test_header_row = next(row for row in range(1, quote_sheet.max_row + 1) if quote_sheet.cell(row, 14).value == "测试费用")
        assert quote_sheet.cell(test_header_row, 15).value == 1500
        assert [quote_sheet.cell(test_header_row + offset, 14).value for offset in range(1, 4)] == [3000, 5000, 10000]
        assert quote_sheet.cell(test_header_row + 1, 16).value == f"=O{test_header_row + 1}*D{quote_rows['报价（MOQ3K）'] - 2}"
        function_row = next(
            row for row in range(1, quote_sheet.max_row + 1)
            if str(quote_sheet.cell(row, 14).value or "").startswith("功能介绍：")
        )
        assert all(
            label not in packaging_labels
            for label in ("彩盒价格", "报客彩盒", "报客彩盒FSC")
        )
        assert test_header_row == function_row + 9
        carton_detail_row = next(
            row for row in range(1, quote_sheet.max_row + 1)
            if quote_sheet.cell(row, 2).value == "纸箱"
        )
        assert quote_sheet.cell(carton_detail_row, 1).value in (None, "")
        assembly_detail_rows = [
            row
            for row in range(1, quote_sheet.max_row + 1)
            if quote_sheet.cell(row, 2).value == "装配工"
        ]
        assert [
            quote_sheet.cell(row, 3).value
            for row in assembly_detail_rows
        ] == [
            "组装马桶（17人/11h/3000）",
            "组装奶瓶（7人/11h/3000）",
            "组装勺子（5人/11h/6000）",
            "组装瓶子（13人/11h/3000）",
            "组装公仔（13人/11h/3000）",
            "包装公仔（40人/11h/3000）",
        ]
        assert [
            quote_sheet.cell(row, 4).value
            for row in assembly_detail_rows
        ] == [
            "=17*$M$4/3000",
            "=7*$M$4/3000",
            "=5*$M$4/6000",
            "=13*$M$4/3000",
            "=13*$M$4/3000",
            "=40*$M$4/3000",
        ]
        assert all(
            quote_sheet.cell(row, 1).value in (None, "")
            for row in assembly_detail_rows
        )
        tax_header_row = next(
            row for row in range(1, quote_sheet.max_row + 1)
            if quote_sheet.cell(row, 3).value == "人民币外购件成本"
        )
        carton_tax_column = next(
            column for column in range(6, 15)
            if quote_sheet.cell(tax_header_row, column).value == "纸箱类"
        )
        assert quote_sheet.cell(tax_header_row + 1, carton_tax_column).value in (None, "")
        assert quote_sheet.cell(tax_header_row + 3, carton_tax_column).value in (None, "")
        summary_row = next(row for row in range(1, quote_sheet.max_row + 1) if quote_sheet.cell(row, 3).value == "旺季价")
        assert quote_sheet.cell(summary_row + 1, 4).value == f"=D{quote_rows['报价（MOQ10K）']}"
        assert quote_sheet["Q6"].value == 0.035
        assert quote_sheet["Q6"].number_format == "0.00%"
        assert quote_sheet["R6"].value == 10000
        assert all(
            workbook[name].sheet_state == "veryHidden"
            for name in ("电子明细", "车缝明细", "车发明细", "装配明细", "审批与版本")
        )
        assert "排摸表" not in workbook.sheetnames
        assert "外购清单" not in workbook.sheetnames
        assert workbook["工程核价依据-报价明细"].sheet_state == "visible"
        assert workbook["工程核价依据-报价明细"]["A1"].value == "上传源表"
        assert workbook["工程核价依据-报价明细"]["A1"].font.bold is True
        assert workbook["工程核价依据-参数"]["A1"].value == "='工程核价依据-报价明细'!A1"
        electronic_sheet = workbook["电子明细"]
        assert [electronic_sheet.cell(3, column).value for column in range(1, 11)] == [
            "父项", "零件名称", "规格", "用量", "单价RMB", "单价HKD", "金额HKD", "税点%", "备注", "来源",
        ]
        electronic_summary_row = next(
            row for row in range(4, electronic_sheet.max_row + 1)
            if electronic_sheet.cell(row, 1).value == "电子成本汇总"
        )
        assert [electronic_sheet.cell(electronic_summary_row, column).value for column in range(1, 5)] == [
            "电子成本汇总", "RMB", "HKD", "公式口径",
        ]
        assert [workbook["车缝明细"].cell(3, column).value for column in range(1, 13)] == [
            "物料名称", "裁片部位", "供应商", "布料MOQ/Y", "低于MOQ/每色费用 RMB",
            "用量/码", "单价 RMB", "汇率", "成本 HKD", "码点", "价钱 HKD", "备注",
        ]
        assert [workbook["车发明细"].cell(3, column).value for column in range(1, 9)] == [
            "#", "名称", "工艺", "重量(g)", "单价(HKD)", "单位", "备注", "金额(HKD)",
        ]
        assert workbook["审批与版本"]["B5"].value == "最终业务放行与客价交接在 P4 实施"
        workbook.close()

        engineering_download = client.post(
            f"/api/internal-quotes/{quote_id}/engineering-data/export"
        )
        assert engineering_download.status_code == 200, engineering_download.text
        assert engineering_download.content.startswith(b"PK")
        assert (
            engineering_download.headers["x-engineering-template-version"]
            == "internal-quote-engineering-template-v1"
        )
        engineering_workbook = load_workbook(
            BytesIO(engineering_download.content),
            data_only=False,
            read_only=False,
        )
        template_workbook = load_workbook(
            ENGINEERING_WORKBOOK_TEMPLATE_PATH,
            data_only=False,
            read_only=False,
        )
        assert engineering_workbook.sheetnames == ["排摸表", "外购清单"]
        mold_schedule = engineering_workbook["排摸表"]
        mold_template = template_workbook["排摸表"]
        assert mold_schedule["O3"].value in (None, "")
        assert mold_schedule["A1"]._style == mold_template["A1"]._style
        assert mold_schedule["A2"]._style == mold_template["A2"]._style
        assert mold_schedule["A5"]._style == mold_template["A5"]._style
        assert mold_schedule["D5"]._style == mold_template["D5"]._style
        assert mold_schedule.row_dimensions[1].height == mold_template.row_dimensions[1].height
        assert mold_schedule.row_dimensions[5].height == mold_template.row_dimensions[5].height
        assert mold_schedule.column_dimensions["B"].width == mold_template.column_dimensions["B"].width
        assert not mold_schedule._images
        assert mold_schedule["O3"].value in (None, "")
        assert [mold_schedule.cell(row, 2).value for row in range(6, 12)] == [
            "左前枪身（橙色）",
            "右前枪身（橙色）",
            "泵杆",
            "击锤",
            "扣机",
            "配件(7件)",
        ]
        assert all(
            mold_schedule.cell(row, column).value in (None, "")
            for row in range(6, 12)
            for column in (4, 5, 9, 17, 18, 19)
        )
        assert {"A6:A7", "C6:C7", "A8:A11", "C8:C11"} <= {
            str(cell_range) for cell_range in mold_schedule.merged_cells.ranges
        }
        purchase_list = engineering_workbook["外购清单"]
        purchase_template = template_workbook["外购清单"]
        assert purchase_list["A1"]._style == purchase_template["A1"]._style
        assert purchase_list["A5"]._style == purchase_template["A5"]._style
        assert purchase_list["J5"]._style == purchase_template["J5"]._style
        assert purchase_list.row_dimensions[6].height == purchase_template.row_dimensions[6].height
        assert purchase_list.column_dimensions["B"].width == purchase_template.column_dimensions["B"].width
        assert not purchase_list._images
        assert str(purchase_list["G4"].value).startswith("文件编号：             版本:")
        purchase_names = {
            purchase_list.cell(row, 2).value
            for row in range(6, purchase_list.max_row + 1)
        }
        assert {"螺丝", "润滑油", "彩盒", "主纸箱", "主控板", "喇叭"} <= purchase_names
        assert all(
            purchase_list.cell(row, 10).value in (None, "")
            for row in range(6, purchase_list.max_row + 1)
        )
        engineering_workbook.close()
        template_workbook.close()
        assert len(client.get(f"/api/internal-quotes/{quote_id}/exports").json()) == 1

        second_response = client.post(f"/api/internal-quotes/{quote_id}/exports")
        assert second_response.status_code == 201
        second = second_response.json()
        history = client.get(f"/api/internal-quotes/{quote_id}/exports")
        assert history.status_code == 200
        by_id = {item["id"]: item for item in history.json()}
        assert by_id[first["id"]]["status"] == "superseded"
        assert by_id[second["id"]]["status"] == "current"

        with db_module.SessionLocal() as db:
            engineering = db.query(quote_models.InternalQuoteSection).filter_by(
                quote_id=quote_id,
                department="engineering",
            ).one()
            engineering.revision += 1
            db.commit()
        refreshed = client.get(f"/api/internal-quotes/{quote_id}/exports").json()
        assert all(item["status"] == "superseded" for item in refreshed)

        logout(client)
        login(client, "iq_p3_export_forbidden", "engineer", "engineering")
        history_only = client.get(f"/api/internal-quotes/{quote_id}/exports")
        assert history_only.status_code == 200
        assert len(history_only.json()) == 2
        forbidden_download = client.get(
            f"/api/internal-quotes/{quote_id}/exports/{second['id']}/download"
        )
        assert forbidden_download.status_code == 403
