"""Warehouse facts are posted once; ambiguous matching stays a business choice."""
from copy import deepcopy
from io import BytesIO, StringIO
import csv
from uuid import uuid4

import openpyxl
import pytest
from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session
from test_indonesia_materials import client, actor, command, material, product, order, supply, allocation, historical_allocation, workspace, BASE
from app.models.indonesia_materials import IndonesiaMaterialRecord as Record, IndonesiaMaterialCommand as Receipt
from app.services.indonesia_warehouse_import import parse_warehouse_receipts


def source(**changes):
    return dict(warehouse='塑胶仓', receipt_no='WH-001', receipt_line='1', actual_date='2026-10-22', material_code='00001',
                material_name='合成螺丝', specification='TEST', unit='PCS', normal_quantity='90', spare_quantity='10',
                purchase_no='PO-WH', purchase_line='0001', purchase_date='2026-10-01', purchase_quantity='100',
                supplier='合成供应商', price='.1', currency='USD') | changes


def content(rows):
    stream = StringIO()
    fields = list(dict.fromkeys(key for row in rows for key in row))
    writer = csv.DictWriter(stream, fieldnames=fields)
    writer.writeheader(); writer.writerows(rows)
    return stream.getvalue().encode('utf-8-sig')


def preview(client, *rows, status=200, **form):
    result = client.post(BASE + '/warehouse-imports/preview', data=dict(factory_id='huaxing', customer='合成客户', recipient='RRI', warehouse='') | form,
                         files={'file': ('real-receipt.csv', content(rows or [source()]), 'text/csv')})
    assert result.status_code == status, result.text
    return result.json()['record'] if status == 200 else result


def choose(preview, row=0, material_id=None, targets=None, **changes):
    record = preview['data']['rows'][row]
    material = next(item for item in record['materials'] if material_id is None or item['id'] == material_id)
    target = record['targets'][material['id']][0]
    return dict(index=record['index'], material_id=material['id'], material_version=material['version'],
        allocations=targets or [dict(order_id=target['order_id'], demand_id=target['demand_id'], order_version=target['order_version'],
                                    normal_quantity=record['source']['normal_quantity'], spare_quantity=record['source']['spare_quantity'])], **changes)


def confirm(client, preview, selections=None, **kwargs):
    return command(client, 'warehouse_import_confirm', dict(shipment_date='2026-10-23', rows=selections or [choose(preview)]), preview, **kwargs)


def catalog(client, quantity='100'):
    m = material(client); product(client, m); o = order(client, quantity=quantity)
    return m, o


def test_no_code_exact_match_posts_real_receipt_allocation_and_unpacked_shipment_once(client):
    m, o = catalog(client)
    client.materials_auth['user'] = actor('read', 'write', 'export')  # No manual Shipping permission.
    imported = preview(client, source(material_code=''))
    assert imported['data']['rows'][0]['auto_material_id'] == m['id']
    key = uuid4().hex
    posted = confirm(client, imported, operation_id=key)
    assert confirm(client, imported, operation_id=key) == posted
    data = workspace(client)
    supplies = [r for r in data['records'] if r['kind'] == 'supply']
    allocations = [r for r in data['records'] if r['kind'] == 'allocation']
    shipments = [r for r in data['records'] if r['kind'] == 'shipment']
    assert len(supplies) == len(allocations) == len(shipments) == 1
    assert supplies[0]['data']['status'] == 'received'
    assert supplies[0]['data']['receipt_confirmation_source'] == 'sales_warehouse_document'
    assert supplies[0]['data']['warehouse_source']['receipt_no'] == 'WH-001'
    assert data['board'][0]['status'] == 'complete'
    shipment = shipments[0]
    assert shipment['data']['status'] == 'draft' and shipment['data']['packages'] == []
    assert not shipment['data'].get('export_asset_id')
    command(client, 'shipment_confirm', {}, shipment, status=422)
    line = shipment['data']['lines'][0]
    payload = {k: shipment['data'][k] for k in ('customer', 'recipient', 'shipment_date', 'lines', 'packages', 'remarks')}
    payload['packages'] = [dict(label='仓库装箱', cartons=1, length_cm='27', width_cm='14.5', height_cm='27.3', carton_weight_kg='.02', items=[line])]
    saved = command(client, 'shipment_save', payload, shipment)
    assert saved['data']['warehouse_import_id'] == imported['id']
    done = command(client, 'shipment_confirm', {}, saved)
    assert done['data']['status'] == 'confirmed'
    assert client.get(BASE + f"/shipments/{done['id']}/export", params={'factory_id': 'huaxing'}).status_code == 200


def test_same_name_and_spec_never_picks_an_arbitrary_material_or_contract(client):
    m, first = catalog(client)
    second = order(client, contract='TEST-B')
    other = material(client, code='00002')
    imported = preview(client, source(material_code=''))
    row = imported['data']['rows'][0]
    assert row['auto_material_id'] is None and len(row['materials']) == 2
    assert len(row['targets'][m['id']]) == 2
    target = next(item for item in row['targets'][m['id']] if item['order_id'] == second['id'])
    picked = choose(imported, material_id=m['id'], targets=[{k: target[k] for k in ('order_id', 'demand_id', 'order_version')} | dict(normal_quantity='90', spare_quantity='10')])
    confirm(client, imported, [picked])
    actual = next(item for item in workspace(client)['records'] if item['kind'] == 'allocation')
    assert actual['data']['order_id'] == second['id']


def test_changed_file_or_case_cannot_reimport_same_warehouse_voucher_line(client):
    catalog(client)
    imported = preview(client)
    confirm(client, imported)
    duplicate = preview(client, source(receipt_no='wh-001', receipt_line='01', normal_quantity='89', spare_quantity='11'))
    assert duplicate['data']['rows'][0]['blocked']
    assert '已经导入' in ' '.join(duplicate['data']['rows'][0]['issues'])
    confirm(client, duplicate, status=409)
    assert len([r for r in workspace(client)['records'] if r['kind'] == 'supply']) == 1


def test_second_row_exceeding_contract_rolls_back_every_post_and_receipt(client):
    catalog(client)
    imported = preview(client, source(normal_quantity='60', spare_quantity='0'), source(receipt_line='2', purchase_line='0002', normal_quantity='60', spare_quantity='0'))
    confirm(client, imported, [choose(imported, 0), choose(imported, 1)], status=409)
    data = workspace(client)
    assert not [r for r in data['records'] if r['kind'] in {'supply', 'allocation', 'shipment'}]
    assert next(r for r in data['records'] if r['id'] == imported['id'])['version'] == imported['version']


def test_one_receipt_can_be_split_between_contracts_with_separate_normal_and_spare_totals(client):
    catalog(client)
    order(client, contract='TEST-B')
    imported = preview(client)
    row = imported['data']['rows'][0]
    targets = row['targets'][row['auto_material_id']]
    plans = [{k: target[k] for k in ('order_id', 'demand_id', 'order_version')} | dict(normal_quantity='45', spare_quantity='5') for target in targets]
    confirm(client, imported, [choose(imported, targets=plans)])
    shipment = next(r for r in workspace(client)['records'] if r['kind'] == 'shipment')
    assert len(shipment['data']['lines']) == 2


@pytest.mark.parametrize('damage', ['material_version', 'order_version', 'normal_quantity', 'foreign_material', 'foreign_customer'])
def test_stale_or_misbound_selections_never_post_receipt(client, damage):
    m, o = catalog(client)
    imported = preview(client)
    choice = choose(imported)
    if damage == 'material_version': choice['material_version'] += 1
    elif damage == 'order_version': choice['allocations'][0]['order_version'] += 1
    elif damage == 'normal_quantity': choice['allocations'][0]['normal_quantity'] = '91'
    elif damage == 'foreign_material': choice['material_id'] = 'not-in-this-factory'
    else:
        product(client, m, customer='另一客户', code='OTHER')
        foreign = order(client, customer='另一客户', item_no='OTHER', contract='OTHER')
        choice['allocations'][0].update(order_id=foreign['id'], demand_id=foreign['data']['demands'][0]['id'], order_version=foreign['version'])
    confirm(client, imported, [choice], status=409 if damage in {'material_version', 'order_version'} else 404 if damage == 'foreign_material' else 422)
    assert not [r for r in workspace(client)['records'] if r['kind'] in {'supply', 'allocation', 'shipment'}]


def test_existing_expected_purchase_and_exact_allocation_are_reused_without_repricing(client):
    m, o = catalog(client)
    prior = supply(client, m)
    assigned = historical_allocation(client, prior, o)
    imported = preview(client, source(purchase_no='PO-TEST'))
    selected = choose(imported, supply_id=prior['id'], supply_version=prior['version'])
    confirm(client, imported, [selected])
    rows = workspace(client)['records']
    assert len([r for r in rows if r['kind'] == 'supply']) == 1
    assert len([r for r in rows if r['kind'] == 'allocation']) == 1
    assert next(r for r in rows if r['id'] == prior['id'])['data']['material_snapshot'] == prior['data']['material_snapshot']
    assert next(r for r in rows if r['kind'] == 'shipment')['data']['lines'][0]['allocation_id'] == assigned['id']


def test_pending_purchase_requires_using_original_instead_of_duplicate_new_source(client):
    m, o = catalog(client)
    supply(client, m)
    imported = preview(client, source(purchase_no='PO-TEST'))
    confirm(client, imported, status=409)
    assert len([r for r in workspace(client)['records'] if r['kind'] == 'supply']) == 1


@pytest.mark.parametrize('change', [dict(price='2'), dict(purchase_quantity='200'), dict(purchase_date='2026-10-02')])
def test_source_purchase_conflict_cannot_be_hidden_by_choosing_existing_batch(client, change):
    m, o = catalog(client)
    prior = supply(client, m)
    imported = preview(client, source(purchase_no='PO-TEST', **change))
    selected = choose(imported, supply_id=prior['id'], supply_version=prior['version'])
    confirm(client, imported, [selected], status=422)
    assert not [r for r in workspace(client)['records'] if r['kind'] == 'shipment']
    assert next(r for r in workspace(client)['records'] if r['id'] == prior['id'])['data']['status'] == 'expected'


@pytest.mark.parametrize('change', [dict(price='￥1.00'), dict(purchase_quantity='1,000'), dict(price='NaN')])
def test_invalid_source_purchase_numbers_are_review_errors_not_unknown_commits(client, change):
    m, o = catalog(client)
    prior = supply(client, m)
    imported = preview(client, source(purchase_no='PO-TEST', **change))
    assert imported['data']['rows'][0]['blocked']
    confirm(client, imported, [choose(imported, supply_id=prior['id'], supply_version=prior['version'])], status=409)
    assert not [r for r in workspace(client)['records'] if r['kind'] == 'shipment']


def test_partial_import_preserves_other_source_and_creates_separate_second_shipment(client):
    catalog(client, quantity='200')
    imported = preview(client, source(), source(receipt_line='2', purchase_line='0002'))
    first = confirm(client, imported, [choose(imported, 0)])
    assert first['data']['status'] == 'partial'
    remaining = next(r for r in workspace(client)['records'] if r['id'] == imported['id'])
    last = confirm(client, remaining, [choose(remaining, 1)])
    assert last['data']['status'] == 'confirmed' and len(last['data']['shipment_ids']) == 2


def test_receiving_import_requires_sales_write_and_rechecks_factory_after_parse(client, monkeypatch):
    catalog(client)
    client.materials_auth['user'] = actor('read', 'receipt_confirm', department='pmc-warehouse')
    preview(client, status=403)
    assert client.get(BASE + '/warehouse-imports/template', params={'factory_id': 'huaxing'}).status_code == 403
    client.materials_auth['user'] = actor('read', 'write')
    monkeypatch.setattr('app.api.indonesia_materials.refresh_user', lambda request, db: actor('read', 'write', factory='huadeng'))
    preview(client, status=403)


def test_missing_or_mismatching_source_values_are_preserved_for_review(client):
    catalog(client)
    imported = preview(client, source(receipt_line=''), source(receipt_no='WH-002', material_code='UNKNOWN'), source(receipt_no='WH-003', unit='KGM'))
    assert imported['data']['rows'][0]['blocked']
    assert len(imported['data']['rows'][1]['materials']) == 1  # Name matches; unknown code still needs business confirmation.
    assert imported['data']['rows'][1]['auto_material_id'] is None
    assert imported['data']['rows'][2]['materials'] == []


def test_explicit_foreign_factory_is_blocked_and_code_name_disagreement_requires_confirmation(client):
    m, o = catalog(client)
    imported = preview(client, source(factory='华登'), source(receipt_no='WH-002', material_name='原单其他名称'))
    assert imported['data']['rows'][0]['blocked']
    assert '厂区' in ' '.join(imported['data']['rows'][0]['issues'])
    assert imported['data']['rows'][1]['auto_material_id'] is None
    assert imported['data']['rows'][1]['materials'][0]['id'] == m['id']


@pytest.mark.parametrize('before_preview', [True, False])
def test_candidate_growth_never_blocks_the_factory_workspace(client, before_preview):
    m, o = catalog(client)
    imported = None if before_preview else preview(client)
    with Session(client.engine) as db:
        for index in range(201):
            payload = deepcopy(o['data']); payload['contract_no'] = f'CAPACITY-{index}'
            db.add(Record(id=uuid4().hex, factory_id='huaxing', kind='order', business_key=uuid4().hex, version=1, data=payload, updated_at='2026-10-22'))
        db.commit()
    if before_preview:
        imported = preview(client)
    data = workspace(client)
    row = next(item for item in data['records'] if item['id'] == imported['id'])['data']['rows'][0]
    assert row['blocked'] and '候选过多' in ' '.join(row['issues'])
    assert len(data['board']) == 202


def test_complete_packing_carries_to_draft_while_repeated_labels_stay_pending(client):
    catalog(client, quantity='200')
    packing = dict(package_label='混箱1', cartons='2', length_cm='27', width_cm='14.5', height_cm='27.3', carton_weight_kg='.02')
    imported = preview(client, source(**packing))
    posted = confirm(client, imported)
    shipment = next(r for r in workspace(client)['records'] if r['id'] in posted['data']['shipment_ids'])
    assert len(shipment['data']['packages']) == 1
    assert shipment['data']['packages'][0]['items'][0]['normal_quantity'] == '90'


@pytest.mark.parametrize('second_packing', [dict(), dict(cartons='')])
def test_shared_box_labels_include_unselected_and_incomplete_source_rows(client, second_packing):
    catalog(client, quantity='200')
    packing = dict(package_label='混箱1', cartons='2', length_cm='27', width_cm='14.5', height_cm='27.3', carton_weight_kg='.02')
    imported = preview(client, source(**packing), source(**(packing | second_packing), receipt_line='2', purchase_line='0002'))
    posted = confirm(client, imported, [choose(imported, 0)])
    shipment = next(r for r in workspace(client)['records'] if r['id'] in posted['data']['shipment_ids'])
    assert shipment['data']['packages'] == [] and shipment['data']['packing_pending']


@pytest.mark.parametrize('value', ['-1', 'NaN', '1.000000001', '1000000000001', '=1+1'])
def test_invalid_source_quantities_block_row_without_erasing_evidence(value):
    rows = parse_warehouse_receipts(content([source(normal_quantity=value)]), 'warehouse.csv')
    assert rows[0]['source_issues'] and rows[0]['raw_values']


def test_total_received_includes_spares_and_duplicate_headers_are_rejected():
    row = source(); row.pop('normal_quantity'); row.update(quantity='100')
    parsed = parse_warehouse_receipts(content([row]), 'warehouse.csv')[0]
    assert parsed['source']['normal_quantity'] == '90'
    with pytest.raises(HTTPException, match='关键表头重复'):
        parse_warehouse_receipts('物料名称,实收数量,实收数量\n螺丝,1,1'.encode(), 'warehouse.csv')


def test_xlsx_preserves_zero_padded_material_codes_and_detects_formulas():
    book = openpyxl.Workbook(); sheet = book.active
    sheet.append(['收料单号', '行号', '实收日期', '物料编码', '实收数量'])
    sheet.append(['WH-001', 1, '2026-10-22', 1, '=1+1']); sheet['D2'].number_format = '00000'
    stream = BytesIO(); book.save(stream); book.close()
    parsed = parse_warehouse_receipts(stream.getvalue(), 'warehouse.xlsx')[0]
    assert parsed['source']['material_code'] == '00001' and parsed['source_issues']


@pytest.mark.parametrize('columns', [5, 81])
def test_xlsx_without_optional_dimensions_still_reads_and_enforces_column_limit(columns):
    from zipfile import ZipFile
    from lxml import etree
    book = openpyxl.Workbook(); sheet = book.active
    sheet.append(['收料单号', '行号', '实收日期', '物料名称', '实收数量'])
    sheet.append(['WH-001', 1, '2026-10-22', '合成螺丝', 10])
    sheet.cell(2, columns, 10)
    stream = BytesIO(); book.save(stream)
    output = BytesIO()
    with ZipFile(BytesIO(stream.getvalue())) as source_zip, ZipFile(output, 'w') as target:
        for info in source_zip.infolist():
            data = source_zip.read(info.filename)
            if info.filename == 'xl/worksheets/sheet1.xml':
                tree = etree.fromstring(data)
                for element in tree.findall('{*}dimension'):
                    tree.remove(element)
                data = etree.tostring(tree)
            target.writestr(info, data)
    if columns == 81:
        with pytest.raises(HTTPException, match='80 列'):
            parse_warehouse_receipts(output.getvalue(), 'warehouse.xlsx')
    else:
        rows = parse_warehouse_receipts(output.getvalue(), 'warehouse.xlsx')
        assert len(rows) == 1 and rows[0]['source']['normal_quantity'] == '10'


def test_real_bounded_warehouse_worker_preserves_rows_and_returns_known_format_errors(monkeypatch, tmp_path):
    import asyncio
    from app.services import customer_order_jobs as jobs
    monkeypatch.setattr(jobs, 'LOCK_PATH', tmp_path / 'warehouse-worker.lock')
    original = content([source()])
    assert asyncio.run(jobs.run_order_job(parse_warehouse_receipts, original, 'warehouse.csv')) == parse_warehouse_receipts(original, 'warehouse.csv')
    with pytest.raises(HTTPException) as error:
        asyncio.run(jobs.run_order_job(parse_warehouse_receipts, b'unknown headers', 'warehouse.csv'))
    assert error.value.status_code == 400 and '未找到实收表头' in str(error.value.detail)
