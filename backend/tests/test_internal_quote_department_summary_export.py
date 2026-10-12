import json
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.services.internal_quote import _rr2_cost_summary
from app.services.internal_quote_calculator import calculate_section
from app.services.internal_quote_excel import _build_summary_sheet
from test_internal_quote_excel_formula_links import _section
from test_internal_quote_purchase_export import arithmetic_value


def build_export(electronic, sewing, factory='huaxing', *, electronic_required=True):
    snapshot = {'fx': {'rmb_hkd': '.85', 'hkd_usd': '7.8'}, 'tax_rates': {'sewing_clothes': '.115'}}
    sections = []
    context = {}
    for code, payload in [('electronic', electronic), ('sewing', sewing)]:
        calc = calculate_section(code, payload, snapshot, 'GROUP-EXPORT')
        assert calc['status'] == 'valid', calc.get('warnings')
        section = _section(code, payload=payload, calculation=calc)
        section.is_required = electronic_required if code == 'electronic' else True
        sections.append(section)
        context[f'{code}_hkd'] = Decimal(calc['totals']['total_hkd']) if section.is_required else Decimal(0)
    context['factory_price_hkd'] = sum(context.values())
    sales = _section('sales', payload={'shipping': {'markup_x': '1.2'}})
    sales.is_required = True
    sections.append(sales)
    summary = _rr2_cost_summary(sections, context, snapshot, factory_id=factory)
    before = [section.calculation_json for section in sections]
    summary_before = json.dumps(summary, default=str)
    workbook = Workbook()
    _build_summary_sheet(workbook, SimpleNamespace(product_name='模拟公仔', quote_no='GROUP-QA',
                         customer='普通客', region_code='mainland', remark='模拟数据'),
                         sections, snapshot, summary, context, [])
    assert before == [section.calculation_json for section in sections]
    assert summary_before == json.dumps(summary, default=str)
    return workbook, workbook['报价明细'], context, summary


def category_rows(sheet, category, visible=True):
    return [row for row in range(1, sheet.max_row + 1)
            if sheet.cell(row, 2).value == category and (not visible or not sheet.row_dimensions[row].hidden)]


@pytest.mark.parametrize('grouped', [False, True])
@pytest.mark.parametrize('detached', [False, True])
def test_electronic_total_and_each_doll_are_exported_once(grouped, detached):
    electronic = {'components': [{'item': 'IC', 'quantity': 2, 'unit_price_rmb': '.85'},
                                 {'item': 'LED', 'quantity': 3, 'unit_price_rmb': '.17'},
                                 {'item': 'Motor 马达', 'quantity': 1, 'unit_price_rmb': '.425'}],
                  'bonding_rmb': '.17', 'labor_rmb': '.255', 'profit_rate_percent': 10}
    if grouped:
        electronic = {'quote_groups': [{**electronic, 'id': 'a', 'name': '主板'},
                                       {'id': 'b', 'name': '灯板', 'quote_mode': 'quick',
                                        'quick_quotes': [{'item': '灯板总价', 'unit_price_rmb': '.85'}]}]}
    sewing = {'groups': [
        {'name': name, 'category': 'clothes', 'labor_rmb': '.85',
         **({'markup_override': override} if detached else {}),
         'materials': [{'item': '布料', 'usage': 2, 'unit_price_rmb': '.85', 'markup': 1},
                       {'item': '丝线', 'usage': 1, 'unit_price_rmb': '.425', 'markup': 1}]}
        for name, override in [('公仔A', '1.03'), ('公仔B', '1.07')]]}
    workbook, sheet, totals, _ = build_export(electronic, sewing)
    try:
        rows = category_rows(sheet, '电子')
        assert len(rows) == 1
        assert sheet.cell(rows[0], 3).value == '电子'
        assert arithmetic_value(sheet, f'D{rows[0]}') == pytest.approx(float(totals['electronic_hkd']))
        assert not category_rows(sheet, '马达')
        cloth = category_rows(sheet, '车衣')
        assert [sheet.cell(row, 3).value for row in cloth] == ['公仔A', '公仔B']
        assert [arithmetic_value(sheet, f'D{row}') for row in cloth] == pytest.approx([3.5, 3.5])
        assert sum(arithmetic_value(sheet, f'D{row}') for row in category_rows(sheet, '车衣', False)) == pytest.approx(float(totals['sewing_hkd']))
        assert not any(dimension.hidden for dimension in sheet.row_dimensions.values())
    finally:
        workbook.close()


@pytest.mark.parametrize('factory', ['huaxing', 'huakang-c'])
def test_mixed_material_labor_summary_has_no_sewing_tax(factory):
    workbook, sheet, _, summary = build_export(
        {'components': [{'item': 'IC', 'quantity': 1, 'unit_price_hkd': 1}]},
        {'groups': [{'name': '公仔A', 'category': 'clothes', 'labor_rmb': '.85', 'markup_override': '1.03',
                     'materials': [{'item': '布', 'usage': 2, 'unit_price_rmb': '.85', 'markup': 1}]}]}, factory)
    try:
        cloth = category_rows(sheet, '车衣')
        assert len(cloth) == 1
        assert arithmetic_value(sheet, f'D{cloth[0]}') == pytest.approx(3)
        assert sheet.cell(cloth[0], 1).value == ''
        base_row = next(cell.row + 1 for row in sheet for cell in row if cell.value == '人民币外购件成本')
        formula = sheet.cell(base_row, 4).value
        tax = next(row for row in summary['t4'] if row['key'] == 'sewcloth13')
        assert tax['rate_percent'] is None
        assert tax['deduction_hkd'] is None
        assert not formula.endswith('+2')
    finally:
        workbook.close()


def test_same_doll_name_with_distinct_multipliers_is_deducted_from_main_once():
    workbook, sheet, totals, _ = build_export(
        {'components': [{'item': 'IC', 'quantity': 1, 'unit_price_hkd': 1}]},
        {'groups': [{'name': '同款公仔', 'category': 'clothes', 'markup_override': markup,
                     'materials': [{'item': '布', 'usage': usage, 'unit_price_rmb': '.85', 'markup': 1}]}
                    for usage, markup in [(2, '1.03'), (3, '1.07')]]})
    try:
        rows = category_rows(sheet, '车衣')
        assert len(rows) == 2
        assert [arithmetic_value(sheet, f'D{row}') for row in rows] == pytest.approx([2, 3])
        assert sum(arithmetic_value(sheet, f'D{row}') for row in rows) == pytest.approx(float(totals['sewing_hkd']))
        assert not any(dimension.hidden for dimension in sheet.row_dimensions.values())
    finally:
        workbook.close()


def test_quick_sewing_keeps_doll_totals_and_ignores_nonparticipating_electronics():
    workbook, sheet, _, _ = build_export(
        {'components': [{'item': 'Motor', 'quantity': 1, 'unit_price_hkd': 99}]},
        {'quote_mode': 'quick', 'quick_quotes': [
            {'doll_name': '公仔A', 'unit_price_hkd': 2},
            {'doll_name': '公仔B', 'unit_price_hkd': 3, 'markup_override': '1.03'}]},
        electronic_required=False)
    try:
        assert not category_rows(sheet, '电子')
        assert not category_rows(sheet, '马达')
        rows = category_rows(sheet, '车衣')
        assert [sheet.cell(row, 3).value for row in rows] == ['公仔A', '公仔B']
        assert [arithmetic_value(sheet, f'D{row}') for row in rows] == pytest.approx([2, 3])
    finally:
        workbook.close()
