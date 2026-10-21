from copy import copy
from datetime import date, datetime
from io import BytesIO
from collections import Counter
import re
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils.datetime import to_excel

from app.services.huakang_a_order_legacy import inspection_archive as archive
from app.services.huakang_a_order_legacy.unified_schedule import organize_360_schedule
from app.services.customer_order_unified import CustomerOrderUnifiedError


TODAY = date(2026, 10, 16)


def template():
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name in archive.SHEETS:
        sheet = book.create_sheet(name)
        ref, item, inspection = (5, 8, 25) if name == 'ITEM表' else (4, 7, 18)
        for row, value in ((4, date(2026, 10, 15)), (5, TODAY),
                           (6, date(2026, 11, 1)), (7, None), (8, '无需验货')):
            sheet.cell(row, ref, f'RL-{row}')
            sheet.cell(row, item, 'ITEM.06')
            sheet.cell(row, inspection, value)
            sheet.cell(row, 40, f'人工备注-{name}-{row}')
        sheet.cell(9, ref, 'RL-9')
        sheet.cell(9, item, 'ITEM.34')
        sheet.cell(9, inspection, '=Z9-7' if name == 'ITEM表' else "='ITEM表'!Y9")
        if name == 'ITEM表':
            sheet['Z9'] = date(2026, 10, 20)
        sheet.cell(10, ref, 'SIM')
        sheet.cell(10, item, 'SIM')
        sheet.cell(10, inspection, date(2020, 1, 1))
        sheet.cell(10, 40, '模拟数据，不参与归档')
        sheet['A20'] = '取消单'
        sheet['B21'] = '取消订单原位不动'
        sheet['A30'] = '已走货订单'
        sheet['B31'] = '已走货原位不动'
        sheet['A4'].font = Font(color='FF008000', bold=True)
        sheet['A4'].fill = PatternFill('solid', fgColor='FFCCFFFF')
        sheet['A4'].comment = Comment('保留订单批注', '业务')
        sheet['AN4'].hyperlink = 'https://example.invalid/order'
        sheet.row_dimensions[4].height = 45
        sheet.merge_cells('AJ4:AK4')
        sheet['AJ4'] = '单行合并备注'
        sheet.conditional_formatting.add('A4:AN15', FormulaRule(formula=['$AN4<>""'], font=Font(color='FFFF0000')))
        validation = DataValidation(type='list', formula1='"A,B"')
        validation.add('AI4:AI15')
        sheet.add_data_validation(validation)
        sheet.print_area = 'A1:AN31'
    book['ITEM表']['K4'] = 24
    book['ITEM表']['L4'] = 6
    book['ITEM表']['M4'] = '=K4/L4'
    book['接单表']['J4'] = "='ITEM表'!$K$4"
    book['接单表']['N4'] = 12.34
    book['接单表']['O4'] = '=J4*N4'
    book['正单评审表']['N4'] = "='接单表'!N4"
    return book


def row_for(sheet, reference):
    col = 5 if sheet.title == 'ITEM表' else 4
    return next(c.row for c in sheet._cells.values() if c.column == col and c.value == reference)


def identities(book):
    return {s.title: Counter(archive._identity(s, r) for r in archive._rows(s, 4, s.max_row + 1)
                            if all(archive._identity(s, r))) for s in book}


def reload(book):
    output = BytesIO()
    book.save(output)
    return openpyxl.load_workbook(BytesIO(output.getvalue()))


def test_moves_only_past_dates_and_keeps_whole_rows_cross_sheet_links_and_regions():
    book = template()
    before = identities(book)
    result = organize_360_schedule(book, today=TODAY)
    assert result['as_of_date'] == '2026-10-16'
    assert all(s['moved_to_completed'] == 2 for s in result['sheets'].values())
    assert identities(book) == before
    for sheet in book:
        past = row_for(sheet, 'RL-4')
        assert past > 31 and row_for(sheet, 'RL-9') > 31
        assert [row_for(sheet, f'RL-{i}') for i in (5, 6, 7, 8)] == [5, 6, 7, 8]
        assert row_for(sheet, 'SIM') == 10
        assert sheet['A20'].value == '取消单' and sheet['B21'].value == '取消订单原位不动'
        assert sheet['A30'].value == '已走货订单' and sheet['B31'].value == '已走货原位不动'
        assert sheet.cell(past, 40).value == f'人工备注-{sheet.title}-4'
        assert sheet.cell(past, 1).font.color.rgb == 'FF008000'
        assert sheet.cell(past, 1).comment.text == '保留订单批注'
        assert sheet.cell(past, 40).hyperlink.ref == f'AN{past}'
        assert sheet.row_dimensions[past].height == 45
        assert f'AJ{past}:AK{past}' in {str(r) for r in sheet.merged_cells.ranges}
        assert any(f'AI{past}' in validation.sqref for validation in sheet.data_validations.dataValidation)
        assert any(f'A{past}' in cf.sqref for cf in sheet.conditional_formatting)
    item, order, review = (row_for(book[n], 'RL-4') for n in archive.SHEETS)
    assert book['ITEM表'].cell(item, 13).value == f'=K{item}/L{item}'
    assert book['接单表'].cell(order, 10).value == f"='ITEM表'!$K${item}"
    assert book['接单表'].cell(order, 15).value == f'=J{order}*N{order}'
    assert book['正单评审表'].cell(review, 14).value == f"='接单表'!N{order}"
    saved = reload(book)
    assert identities(saved) == before
    again = organize_360_schedule(saved, today=TODAY)
    assert all(not s['moved_to_completed'] and not s['moved_to_pending'] for s in again['sheets'].values())
    assert identities(saved) == before


def test_edited_completed_date_returns_to_pending_without_duplicate_or_shipping_changes():
    book = template()
    organize_360_schedule(book, today=TODAY)
    before = identities(book)
    for sheet in book:
        row = row_for(sheet, 'RL-4')
        sheet.cell(row, 25 if sheet.title == 'ITEM表' else 18, TODAY)
    result = organize_360_schedule(book, today=TODAY)
    assert all(s['moved_to_pending'] == 1 for s in result['sheets'].values())
    assert identities(book) == before
    assert all(row_for(s, 'RL-4') < 20 for s in book)


def test_excel_serial_dates_use_workbook_epoch_and_today_is_not_completed():
    book = template()
    from openpyxl.utils.datetime import CALENDAR_MAC_1904
    book.epoch = CALENDAR_MAC_1904
    for sheet in book:
        col = 25 if sheet.title == 'ITEM表' else 18
        sheet.cell(4, col, to_excel(datetime(2026, 10, 15), book.epoch))
        sheet.cell(5, col, to_excel(datetime(2026, 10, 16), book.epoch))
    organize_360_schedule(book, today=TODAY)
    assert all(row_for(s, 'RL-4') > 31 and row_for(s, 'RL-5') == 5 for s in book)


def test_summary_missing_row_is_not_fabricated_and_ambiguous_dates_do_not_move():
    book = template()
    summary = book['接单表']
    summary['D4'] = 'OTHER'
    summary['R4'] = '=UNKNOWN()'
    item = book['ITEM表']
    item['E11'], item['H11'], item['Y11'] = 'RL-9', 'ITEM.34', date(2026, 12, 1)
    summary['R9'] = '=UNKNOWN()'
    before = identities(book)
    result = organize_360_schedule(book, today=TODAY)
    assert identities(book) == before
    assert row_for(summary, 'RL-9') == 9
    assert any('不同验货日期' in w for w in result['warnings'])


def test_unsafe_partial_range_is_refused_before_any_cell_moves():
    book = template()
    book['ITEM表']['A1'] = '=SUM(K4:K6)'
    before = identities(book)
    with pytest.raises(CustomerOrderUnifiedError, match='不能保持连续引用'):
        organize_360_schedule(book, today=TODAY)
    assert identities(book) == before
    assert all(row_for(s, 'RL-4') == 4 for s in book)
    assert all(not any(c.value == archive.COMPLETED for c in s._cells.values()) for s in book)


def test_unknown_formula_and_unparseable_dates_remain_pending_with_warning():
    book = template()
    book['ITEM表']['Y4'] = '=UNKNOWN()'
    result = organize_360_schedule(book, today=TODAY)
    assert row_for(book['ITEM表'], 'RL-4') == 4
    assert any('公式无可靠结果' in warning for warning in result['warnings'])


def test_clearing_completed_item_date_returns_linked_summary_rows_to_pending():
    book = template()
    for name in ('接单表', '正单评审表'):
        book[name]['R4'] = '=IF(LEN(\'ITEM表\'!Y4)=0,"",\'ITEM表\'!Y4)'
    organize_360_schedule(book, today=TODAY)
    book['ITEM表'].cell(row_for(book['ITEM表'], 'RL-4'), 25).value = None
    result = organize_360_schedule(book, today=TODAY)
    assert all(s['moved_to_pending'] == 1 for s in result['sheets'].values())
    assert all(row_for(sheet, 'RL-4') < 20 for sheet in book)


def test_order_remark_is_not_a_section_marker_and_summary_simulation_is_excluded():
    book = template()
    for sheet in book:
        sheet['AN5'] = '已完成'
    for name in ('接单表', '正单评审表'):
        book[name]['AN10'] = "='ITEM表'!AN10"
    result = organize_360_schedule(book, today=TODAY)
    assert all(s['moved_to_completed'] == 2 for s in result['sheets'].values())
    assert all(row_for(sheet, 'SIM') == 10 for sheet in book)


def with_cached_formula(book, formula):
    book['ITEM表']['Y4'] = formula
    book['ITEM表']['Z4'] = date(2026, 10, 20)
    raw = BytesIO()
    book.save(raw)
    output = BytesIO()
    with ZipFile(BytesIO(raw.getvalue())) as before, ZipFile(output, 'w') as after:
        for name in before.namelist():
            content = before.read(name)
            if name == 'xl/worksheets/sheet1.xml':
                root = ET.fromstring(content)
                cell = root.find('.//m:c[@r="Y4"]', archive.NS)
                cell.find('m:v', archive.NS).text = str(to_excel(datetime(2026, 10, 13)))
                content = ET.tostring(root)
            after.writestr(name, content)
    return output.getvalue()


@pytest.mark.parametrize('formula', ['=Z4-7', '=IF(LEN(Z4)=0,"",Z4-7)'])
def test_cached_date_cannot_survive_changed_formula_inputs(formula):
    payload = with_cached_formula(template(), formula)
    book = openpyxl.load_workbook(BytesIO(payload))
    cached = archive._cache(payload)
    assert archive._date_value(book, book['ITEM表'], 'Y4', cached) == date(2026, 10, 13)
    book['ITEM表']['Z4'] = date(2026, 11, 20)
    value = archive._date_value(book, book['ITEM表'], 'Y4', cached)
    assert value == (date(2026, 11, 13) if formula == '=Z4-7' else None)


def lookup_formula(row):
    return (f'=INDEX(ITEM表!$Y$4:$Y$100,SUMPRODUCT((ITEM表!$E$4:$E$100=$D{row})*'
            f'(ITEM表!$H$4:$H$100=$G{row})*(ROW(ITEM表!$E$4:$E$100)-ROW(ITEM表!$E$4)+1)))')


def calculate_lookup(book, sheet, row):
    """Independent arithmetic for the fixture's actual Excel INDEX expression."""
    formula = sheet.cell(row, 18).value
    pattern = (r'=INDEX\(ITEM表!\$Y\$(\d+):\$Y\$(\d+),SUMPRODUCT\(\(ITEM表!\$E\$(\d+):\$E\$(\d+)=\$D(\d+)\)\*'
               r'\(ITEM表!\$H\$(\d+):\$H\$(\d+)=\$G(\d+)\)\*\(ROW\(ITEM表!\$E\$(\d+):\$E\$(\d+)\)-ROW\(ITEM表!\$E\$(\d+)\)\+1\)\)\)')
    match = re.fullmatch(pattern, formula)
    assert match, formula
    y_first, y_last, e_first, e_last, d_row, h_first, h_last, g_row, rows_first, rows_last, baseline = map(int, match.groups())
    assert (e_first, e_last) == (h_first, h_last) == (rows_first, rows_last)
    item = book['ITEM表']
    reference, product = sheet.cell(d_row, 4).value, sheet.cell(g_row, 7).value
    index = sum((r - baseline + 1) for r in range(rows_first, rows_last + 1)
                if getattr(item._cells.get((r, 5)), 'value', None) == reference
                and getattr(item._cells.get((r, 8)), 'value', None) == product)
    assert 1 <= index <= y_last - y_first + 1
    return item.cell(y_first + index - 1, 25).value


def test_index_sumproduct_results_survive_moving_the_lookup_range_first_order():
    book = template()
    for sheet_name in ('接单表', '正单评审表'):
        for row in (4, 5):
            book[sheet_name].cell(row, 18, lookup_formula(row))
            assert calculate_lookup(book, book[sheet_name], row) == (date(2026, 10, 15) if row == 4 else TODAY)
    organize_360_schedule(book, today=TODAY)
    for phase in (book, reload(book)):
        for sheet_name in ('接单表', '正单评审表'):
            for reference, expected in [('RL-4', date(2026, 10, 15)), ('RL-5', TODAY)]:
                row = row_for(phase[sheet_name], reference)
                value = calculate_lookup(phase, phase[sheet_name], row)
                assert (value.date() if isinstance(value, datetime) else value) == expected
        result = organize_360_schedule(phase, today=TODAY)
        assert all(s['moved_to_completed'] == s['moved_to_pending'] == 0 for s in result['sheets'].values())


def test_original_review_array_price_lookup_allows_a_different_column_row_origin():
    book = template()
    formula = ('=IF(OR(LEN($D4)=0,LEN($G4)=0,COUNTIFS(接单表!$D$4:$D$1062,$D4,接单表!$G$4:$G$1062,$G4)<>1,'
               'COUNTIFS(接单表!$D$4:$D$1062,$D4,接单表!$G$4:$G$1062,$G4,接单表!$N$4:$N$1062,"<>")<>1),"",'
               'IFERROR(_xlfn.SINGLE(INDEX(接单表!$N$4:$N$1062,SUMPRODUCT((接单表!$D$4:$D$1062=$D4)*'
               '(接单表!$G$4:$G$1062=$G4)*(ROW(接单表!$D$4:$D$1062)-ROW(接单表!$C$4)+1)))),""))')
    book['正单评审表']['N4'] = ArrayFormula(ref='N4', text=formula)
    organize_360_schedule(book, today=TODAY)
    for phase in (book, reload(book)):
        review = phase['正单评审表']
        review_row = row_for(review, 'RL-4')
        cell = review.cell(review_row, 14)
        assert isinstance(cell.value, ArrayFormula) and cell.value.ref == f'N{review_row}'
        text = cell.value.text
        first = int(re.search(r'INDEX\(接单表!\$N\$(\d+)', text).group(1))
        baseline = int(re.search(r'-ROW\(接单表!\$C\$(\d+)\)', text).group(1))
        order = phase['接单表']
        matches = [r for r in archive._rows(order, 4, 1063) if archive._identity(order, r) == archive._identity(review, review_row)]
        index = sum(r - baseline + 1 for r in matches)
        assert len(matches) == 1 and first + index - 1 == matches[0]
        assert order.cell(first + index - 1, 14).value == 12.34


@pytest.mark.parametrize('formula', ['=ROW(ITEM表!$E$4)', '=SUMPRODUCT(ROW(ITEM表!$E$4:$E$100))'])
def test_unproven_row_formula_is_refused_before_order_movement(formula):
    book = template()
    book['接单表']['A1'] = formula
    with pytest.raises(CustomerOrderUnifiedError, match='ROW 公式'):
        organize_360_schedule(book, today=TODAY)
    assert all(row_for(sheet, 'RL-4') == 4 for sheet in book)


def test_global_and_sheet_local_names_follow_the_order_and_survive_reopen():
    book = template()
    book.defined_names.add(DefinedName('GlobalQuantity', attr_text="'ITEM表'!$K$4"))
    book['ITEM表'].defined_names.add(DefinedName('OrderQuantity', attr_text="'ITEM表'!$K$4"))
    book['ITEM表'].defined_names.add(DefinedName('LocalQuantity', attr_text='$K$4'))
    organize_360_schedule(book, today=TODAY)
    for phase in (book, reload(book)):
        row = row_for(phase['ITEM表'], 'RL-4')
        assert phase.defined_names['GlobalQuantity'].attr_text == f"'ITEM表'!$K${row}"
        assert phase['ITEM表'].defined_names['OrderQuantity'].attr_text == f"'ITEM表'!$K${row}"
        local = phase['ITEM表'].defined_names['LocalQuantity'].attr_text
        assert phase['ITEM表'][local.replace('$', '')].value == 24


def test_multiarea_validation_uses_one_original_relative_formula_anchor():
    book = template()
    item = book['ITEM表']
    item['A4'], item['A10'] = 1, -1
    validation = DataValidation(type='custom', formula1='A4>0')
    validation.add('AO4:AO6')
    validation.add('AO10:AO12')
    item.add_data_validation(validation)
    organize_360_schedule(book, today=TODAY)
    for phase in (book, reload(book)):
        sheet = phase['ITEM表']
        moved = row_for(sheet, 'RL-4')
        for target, expected in [(f'AO{moved}', True), ('AO10', False)]:
            rule = next(v for v in sheet.data_validations.dataValidation if target in v.sqref)
            match = re.fullmatch(r'(A\d+)>0', rule.formula1)
            assert match and (sheet[match.group(1)].value > 0) == expected


def test_conflicting_dates_across_unique_three_sheet_identity_stay_pending():
    book = template()
    book['接单表']['R4'] = date(2026, 12, 1)
    result = organize_360_schedule(book, today=TODAY)
    assert all(row_for(sheet, 'RL-4') == 4 for sheet in book)
    assert any('验货日期存在冲突' in w for w in result['warnings'])
    book['接单表']['R4'] = date(2026, 10, 15)
    organize_360_schedule(book, today=TODAY)
    book['接单表'].cell(row_for(book['接单表'], 'RL-4'), 18, date(2026, 12, 1))
    result = organize_360_schedule(book, today=TODAY)
    assert all(s['moved_to_pending'] == 1 for s in result['sheets'].values())
    assert all(row_for(sheet, 'RL-4') < 20 for sheet in book)


def test_real_multiple_batches_are_not_conflicts_and_protected_history_is_excluded():
    book = template()
    for sheet in book:
        reference, product, inspection = (5, 8, 25) if sheet.title == 'ITEM表' else (4, 7, 18)
        sheet.cell(5, reference, 'RL-4')  # Same identity, distinct actual delivery batch.
        sheet.cell(21, reference, 'RL-4')
        sheet.cell(21, product, 'ITEM.06')
        sheet.cell(21, inspection, date(2027, 1, 1))
    result = organize_360_schedule(book, today=TODAY)
    assert all(s['moved_to_completed'] == 2 for s in result['sheets'].values())
    assert not any('冲突' in w for w in result['warnings'])
    for sheet in book:
        assert sheet.cell(5, 5 if sheet.title == 'ITEM表' else 4).value == 'RL-4'
        assert sheet.cell(21, 25 if sheet.title == 'ITEM表' else 18).value == date(2027, 1, 1)


def test_unknown_summary_formula_does_not_infer_a_date_from_matching_item():
    book = template()
    book['接单表']['R4'] = '=UNKNOWN()'
    result = organize_360_schedule(book, today=TODAY)
    assert row_for(book['接单表'], 'RL-4') == 4
    assert any('接单表 第4行验货日期公式无可靠结果' in w for w in result['warnings'])
