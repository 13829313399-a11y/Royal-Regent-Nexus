from io import BytesIO

import openpyxl
import pytest

from app.services import customer_order_unified as service
from app.services import customer_order_huaxing_unified as mapping


def template(customer='buzzbee'):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    names = mapping.BUZZBEE_SHEETS if customer == 'buzzbee' else ('ITEM表', mapping.DICKIE_EXTRA) if customer == 'dickie' else ('ITEM表',)
    for i, name in enumerate((*names, '接单表', '正单评审表')):
        ws = wb.create_sheet(name)
        for col, header in enumerate(service.ITEM_HEADERS if name in names else service.ORDER_HEADERS, 1):
            ws.cell(3, col, header)
        ws.cell(8 + i, 1, '取消单')
        ws.cell(9 + i, 2, '保留人工内容')
    return wb


def data(wb):
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_multilane_headers_and_stock_product_lookup_are_not_order_identity():
    wb = template()
    ws = wb[mapping.BUZZBEE_SHEETS[0]]
    ws['C4'], ws['H4'], ws['I4'] = 'PRODUCTION', '40210-1', '转盘枪'
    ws['H12'], ws['D12'], ws['K12'], ws['L12'] = '58208-2', '53102', 1200, '0/24'
    content = data(wb)
    service.ensure_unified_schedule('new.xlsx', content, factory_id='huaxing', customer_code='buzzbee')
    assert service.is_unified_schedule(content)
    history = service.read_unified_history(content, factory_id='huaxing', customer_code='buzzbee')
    assert history[0].reference_no == ''
    assert history[1].reference_no == '53102'
    assert history[1].units_per_carton == '24'
    with pytest.raises(service.CustomerOrderUnifiedError):
        service.ensure_unified_schedule('new.xlsx', content, factory_id='huadeng', customer_code='spin')


def test_summary_links_actual_lane_and_independent_row_preserving_history():
    wb = template()
    first, second = mapping.BUZZBEE_SHEETS[:2]
    wb[first]['H4'], wb[first]['I4'] = '40210-1', '转盘枪'
    wb[second]['H4'], wb[second]['I4'] = '57520-3', '水枪'
    wb['接单表']['B4'] = '不能覆盖的手工备注'
    wb['接单表']['B5'] = '第二条备注'
    rows = [dict(id=str(i), product_no=item, quantity='12', units_per_carton='6', reference_no=f'SO{i}',
                 item_sheet_name=lane, customer_name='客', issues=[], _huaxing_customer='buzzbee')
            for i, (lane, item) in enumerate([(first, '40210-1'), (second, '57520-3')])]
    mapping.export_rows(wb, rows)
    assert wb[first]['E5'].value == 'SO0'
    assert wb[second]['E5'].value == 'SO1'
    assert wb['接单表']['D6'].value == 'SO0'
    assert wb['接单表']['J6'].value == f'=IF(LEN(\'{first}\'!K5)=0,"",\'{first}\'!K5)'
    assert wb['正单评审表']['J5'].value == f'=IF(LEN(\'{second}\'!K5)=0,"",\'{second}\'!K5)'
    assert wb[first]['I4'].value == '转盘枪'
    assert wb['接单表']['B4'].value == '不能覆盖的手工备注'


def test_buzzbee_variant_requires_unique_order_or_customer_evidence():
    history = [service.UnifiedHistoryRow(row=4, product_no='57520-1', customer_name='A', reference_no='OLD', source_sheet='ITEM表-水枪'),
               service.UnifiedHistoryRow(row=7, product_no='57520-3', customer_name='B', reference_no='53102', source_sheet='ITEM表-水枪')]
    row = dict(product_no='57520', customer_name='B', contract_no='53102', issues=[])
    mapping.prepare_row(row, history, 'buzzbee', mapping.BUZZBEE_SHEETS)
    assert row['product_no'] == '57520-3'
    assert row['item_sheet_name'] == 'ITEM表-水枪'
    row = dict(product_no='57520', customer_name='UNKNOWN', contract_no='NEW', issues=[])
    mapping.prepare_row(row, history, 'buzzbee', mapping.BUZZBEE_SHEETS)
    assert any(i['code'] == 'ambiguous_product_variant' for i in row['issues'])


@pytest.mark.parametrize('code,record,expected', [
    ('360', {'production_no': 'RL-450608-3-1', 'contract_no': '450608', 'customer_po': 'CINV-0004380'}, ('', 'RL-450608-3-1', 'CINV-0004380')),
    ('yinhui', {'so_no': '2300000568', 'contract_no': '4500001923'}, ('', '2300000568', '4500001923')),
    ('edu', {'huaxing_po': 'EDUHX00084', 'contract_no': '4500132298/FDSO-26029', 'customer_po': '226403148'}, ('4500132298/FDSO-26029', 'EDUHX00084', '226403148')),
    ('shushupapa', {'po_number': '2026-0402', 'customer_po': '2022821'}, ('2026-0402', '2026-0402', '2022821')),
])
def test_customer_identity_columns_and_no_legacy_display_strings(code, record, expected):
    row = dict(contract_no=record.get('po_number'), standard='PO价格币种USD', packaging='标签/说明书', po_no='WRONG')
    mapping.map_record(row, record, code)
    assert tuple(row[k] for k in ('contract_no', 'reference_no', 'po_no')) == expected
    if code != 'edu':
        assert row['standard'] == ''
        assert row['packaging'] == ''


def test_edu_number_reads_reference_after_cancel_and_reuses_matching_order(monkeypatch):
    from app.services import customer_order_huaxing as legacy
    history = [service.UnifiedHistoryRow(row=20, contract_no='C1', reference_no='EDUHX00082-2', po_no='P1', product_no='I1')]
    assert legacy._next_edu_number([{'huaxing_po': h.reference_no} for h in history]) == 83
    row = dict(contract_no='C1', reference_no='EDUHX00083', po_no='P1', product_no='I1', issues=[])
    mapping.prepare_row(row, history, 'edu', ['ITEM表'])
    assert row['reference_no'] == 'EDUHX00082-2'


def test_unknown_lane_cannot_be_exported_by_confirming_an_unresolved_issue():
    from importlib import import_module
    wb = template()
    with pytest.raises(import_module('app.services.customer_order_unified').CustomerOrderUnifiedError, match='分类页'):
        mapping.export_rows(wb, [dict(item_sheet_name='', product_no='UNKNOWN', _huaxing_customer='buzzbee')])


def test_broken_reserved_review_price_uses_exact_new_order_row_only():
    wb = template('edu')
    wb['正单评审表']['N4'] = '=IFERROR(INDEX(接单表!N4:N30,ROW(接单表!#REF!)),"")'
    wb['正单评审表']['D5'] = 'OLD'
    wb['正单评审表']['N5'] = '=IFERROR(接单表!#REF!,"")'
    wb['正单评审表']['N6'] = '=IFERROR(INDEX(接单表!N4:N30,ROW(接单表!#REF!)),"")'
    row = dict(item_sheet_name='ITEM表', product_no='I1', reference_no='NEW', quantity='12',
               units_per_carton='6', _huaxing_customer='edu')
    mapping.export_rows(wb, [row])
    assert wb['正单评审表']['N6'].value == '=IF(LEN(\'接单表\'!N4)=0,"",\'接单表\'!N4)'
    assert wb['正单评审表']['N5'].value == '=IFERROR(接单表!#REF!,"")'
    assert wb['接单表']['N4'].value is None


def test_silverlit_numeric_so_matches_existing_schedule():
    history = [service.UnifiedHistoryRow(row=4, reference_no='2300000568', po_no='4500001923', product_no='886380NT00101', quantity='3600')]
    row = dict(id='silverlit', quantity='3600', product_no='886380NT00101', issues=[])
    mapping.map_record(row, dict(so_no='SO2300000568', contract_no='PO4500001923'), 'yinhui')
    mapping.prepare_row(row, history, 'yinhui', ['ITEM表'])
    service._enrich_row(row, history, '银辉')
    assert any(i['code'] == 'duplicate_existing_order' for i in row['issues'])


def test_manual_product_change_reselects_lane_and_cannot_reuse_duplicate_confirmation():
    from importlib import import_module
    history = [service.UnifiedHistoryRow(row=4, reference_no='EXISTING', product_no='WATER', source_sheet='ITEM表-水枪')]
    row = dict(id='1', product_no='WATER', reference_no='NEW', contract_no='NEW', po_no='', item_sheet_name='ITEM表-子弹枪')
    original = {'1': ('DART', 'NEW', 'NEW', '')}
    mapping.validate_edited_identities({'rows': [row]}, original, history, mapping.BUZZBEE_SHEETS)
    assert row['item_sheet_name'] == 'ITEM表-水枪'
    row['reference_no'] = 'EXISTING'
    with pytest.raises(import_module('app.services.customer_order_unified').CustomerOrderUnifiedError, match='原重复确认不适用'):
        mapping.validate_edited_identities({'rows': [row]}, original, history, mapping.BUZZBEE_SHEETS)


def test_buzzbee_other_customer_variant_is_not_automatic_product_mapping():
    history = [service.UnifiedHistoryRow(row=4, reference_no='OLD', customer_name='OTHER', product_no='12345-2', source_sheet='ITEM表-水枪')]
    row = dict(product_no='12345', customer_name='NEW', contract_no='NEW', issues=[])
    mapping.prepare_row(row, history, 'buzzbee', mapping.BUZZBEE_SHEETS)
    assert row['product_no'] == '12345'
    assert any(i['code'] == 'ambiguous_product_variant' for i in row['issues'])


def test_360_single_historical_suffix_and_split_release_stays_blocked():
    history = [service.UnifiedHistoryRow(row=4, reference_no='RL-450608-2-1', product_no='2716030120', quantity='600'),
               service.UnifiedHistoryRow(row=5, reference_no='RL-450608-2-2', product_no='2716030120', quantity='72')]
    row = dict(reference_no='RL-450608-2', product_no='2716030120', quantity='672', issues=[])
    mapping.prepare_row(row, history, '360', ['ITEM表'])
    assert any(i['code'] == 'existing_split_release' for i in row['issues'])


def test_360_preserves_full_item_digits_and_reconciles_latest_identity():
    history = [service.UnifiedHistoryRow(row=4, reference_no='RL-450608-3-1', po_no='CINV-0004380', product_no='2716030120')]
    row = dict(product_no='LEGACY', issues=[])
    mapping.map_record(row, dict(item_no='2716030', item_full='2716030120.03', production_no='RL-450608-3', customer_po='0004380'), '360')
    mapping.prepare_row(row, history, '360', ['ITEM表'])
    assert (row['product_no'], row['reference_no'], row['po_no']) == ('2716030120', 'RL-450608-3-1', 'CINV-0004380')


@pytest.mark.parametrize('code', ['360', 'yinhui'])
def test_english_po_description_does_not_fill_chinese_column(code):
    row = dict(product_no='I', product_name_zh='English description', issues=[])
    mapping.map_record(row, dict(product_name='English description'), code)
    assert row['product_name_zh'] == ''
    assert row['product_name_en'] == 'English description'


def test_manual_correction_refreshes_only_history_fields():
    history = [service.UnifiedHistoryRow(row=4, product_no='NEW', units_per_carton='24', product_name_zh='新产品')]
    row = dict(id='1', product_no='NEW', reference_no='R', contract_no='R', po_no='', units_per_carton='6',
               product_name_zh='原产品', product_name_en='Explicit PO description', lineage={'units_per_carton':'唯一历史值', 'product_name_zh':'当前排期'})
    mapping.validate_edited_identities({'rows':[row]}, {'1':('OLD','R','R','')}, history, ['ITEM表'])
    assert row['units_per_carton'] == '24'
    assert row['product_name_zh'] == '新产品'
    assert row['product_name_en'] == 'Explicit PO description'


def test_missing_reference_manual_contract_fallback_is_checked_before_export(monkeypatch):
    from importlib import import_module
    current = import_module('app.services.customer_order_unified')
    wb=template('maxx')
    wb['ITEM表']['E4'],wb['ITEM表']['H4'],wb['ITEM表']['K4']='EXISTING','12345',12
    monkeypatch.setattr(current, '_create_customer_rows', lambda *a: ([dict(id='1',product_no='12345',quantity='12',issues=[])], [], 'TEST'))
    with pytest.raises(current.CustomerOrderUnifiedError, match='原重复确认不适用'):
        current.export_unified_customer_schedule(factory_id='huaxing',customer_code='maxx',received_date='2026-09-08',
            po_files=[('po.pdf',b'po')],schedule_file_name='new.xlsx',schedule_content=data(wb),
            manual_overrides=[dict(row_id='1',field='contract_no',value='EXISTING',reason='人工更正')])
