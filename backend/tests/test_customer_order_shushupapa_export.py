"""Synthetic export cases for Shushupapa detail rows after legacy subtotals."""
from io import BytesIO

import openpyxl
import pytest

from app.services import customer_order_huaxing_unified as mapping
from app.services import customer_order_unified as unified


def workbook():
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name in ('ITEM表', '接单表', '正单评审表'):
        sheet = book.create_sheet(name)
        headers = unified.ITEM_HEADERS if name == 'ITEM表' else unified.ORDER_HEADERS
        for column, header in enumerate(headers, 1):
            sheet.cell(3, column, header)
    item = book['ITEM表']
    item['D4'], item['E4'], item['H4'] = 'SYNTH-OLD', 'SYNTH-OLD', 'SYNTH-ITEM-OLD'
    item['K4'], item['L4'], item['M4'] = 10, 2, '=K4/L4'
    item['A7'], item['D8'] = '取消单', 'SYNTH-CANCELLED'
    for name, total, cancel in (('接单表', 5, 7), ('正单评审表', 6, 8)):
        sheet = book[name]
        sheet['C4'], sheet['D4'], sheet['G4'] = 'SYNTH-OLD', 'SYNTH-OLD', 'SYNTH-ITEM-OLD'
        sheet['J4'], sheet['N4'], sheet['P4'] = 10, 15, 2
        sheet['O4'], sheet['Q4'] = '=J4*N4', '=J4*P4'
        sheet.cell(total, 1, '合计')
        # This is deliberately the last row containing formulas. The legacy
        # slot copier uses it as the template for newly appended details.
        for column in ('N', 'O', 'P', 'Q'):
            sheet[f'{column}{total}'] = f'=SUM({column}4:{column}{total - 1})'
        sheet.cell(cancel, 1, '取消单')
        sheet.cell(cancel + 1, 4, 'SYNTH-CANCELLED')
    extra = book.create_sheet('用户资料')
    extra['A1'], extra['B1'] = '保留此页', '=2+3'
    return book


def order(index=1, **changes):
    row = dict(
        id=f'synthetic-{index}', item_sheet_name='ITEM表',
        product_no=f'SYNTH-ITEM-{index}', contract_no=f'SYNTH-C-{index}',
        reference_no=f'SYNTH-SO-{index}', po_no=f'SYNTH-PO-{index}',
        customer_name='SYNTHETIC CUSTOMER', quantity='120', units_per_carton='12',
        unit_price_hkd='15.25', unit_price_usd='2.50',
        issues=[], _huaxing_customer='shushupapa',
    )
    row.update(changes)
    return row


def find_row(sheet, reference):
    column = 5 if sheet.title == 'ITEM表' else 4
    return next(r for r in range(4, sheet.max_row + 1)
                if sheet.cell(r, column).value == reference)


def amount_formula(number, price):
    return (f'=IF(OR(LEN($J{number})=0,LEN(${price}{number})=0),'
            f'"",ROUND($J{number}*${price}{number},2))')


def reload(book):
    stream = BytesIO()
    book.save(stream)
    return openpyxl.load_workbook(BytesIO(stream.getvalue()))


def assert_detail_prices(book, row):
    order_sheet, review = book['接单表'], book['正单评审表']
    a, b = (find_row(sheet, row['reference_no']) for sheet in (order_sheet, review))
    for column, field in (('N', 'unit_price_hkd'), ('P', 'unit_price_usd')):
        assert order_sheet[f'{column}{a}'].value == unified._number(row.get(field))
        ref = f"'接单表'!{column}{a}"
        assert review[f'{column}{b}'].value == f'=IF(LEN({ref})=0,"",{ref})'
    for sheet, number in ((order_sheet, a), (review, b)):
        assert sheet[f'O{number}'].value == amount_formula(number, 'N')
        assert sheet[f'Q{number}'].value == amount_formula(number, 'P')
        for column in ('N', 'O', 'P', 'Q'):
            assert sheet[f'{column}{number}'].number_format == '#,##0.00'
    return a, b


def test_new_detail_replaces_inherited_totals_and_uses_exact_paired_summary_row():
    book = workbook()
    row = order()
    mapping.export_rows(book, [row])
    a, b = assert_detail_prices(book, row)
    assert (a, b) == (6, 7)
    detail = find_row(book['ITEM表'], row['reference_no'])
    for sheet, number in ((book['接单表'], a), (book['正单评审表'], b)):
        assert sheet[f'J{number}'].value == f'=IF(LEN(\'ITEM表\'!K{detail})=0,"",\'ITEM表\'!K{detail})'


@pytest.mark.parametrize('hkd,usd', [
    (None, None), ('', ''), ('unconfirmed', None),
    (None, '2.50'), ('15.25', None), ('0', '0'), (0, 0),
])
def test_unknown_price_stays_empty_and_zero_is_preserved(hkd, usd):
    book = workbook()
    row = order(unit_price_hkd=hkd, unit_price_usd=usd)
    mapping.export_rows(book, [row])
    assert_detail_prices(reload(book), row)


def test_six_rows_survive_save_reload_without_cross_order_totals_or_price_references():
    book = workbook()
    rows = [order(i, unit_price_hkd=str(i + 10), unit_price_usd=str(i))
            for i in range(1, 7)]
    mapping.export_rows(book, rows)
    reopened = reload(book)
    for row in rows:
        assert_detail_prices(reopened, row)
    for name in ('ITEM表', '接单表', '正单评审表'):
        sheet = reopened[name]
        marker = next(c.row for row in sheet for c in row if c.value == '取消单')
        assert all(find_row(sheet, row['reference_no']) < marker for row in rows)


def test_history_totals_cancelled_content_and_other_sheet_are_retained():
    book = workbook()
    protected = {}
    for name, end in (('ITEM表', 4), ('接单表', 5), ('正单评审表', 6)):
        sheet = book[name]
        protected[name] = {cell.coordinate: (cell.value, tuple(cell._style or (0,) * 9))
                           for row in sheet.iter_rows(max_row=end) for cell in row}
    mapping.export_rows(book, [order(i) for i in range(1, 7)])
    for name, cells in protected.items():
        for coordinate, (value, style) in cells.items():
            assert book[name][coordinate].value == value
            assert tuple(book[name][coordinate]._style or (0,) * 9) == style
        assert sum(cell.value == 'SYNTH-CANCELLED' for row in book[name] for cell in row) == 1
    assert book['用户资料']['A1'].value == '保留此页'
    assert book['用户资料']['B1'].value == '=2+3'


@pytest.mark.parametrize('customer', ['maxx', 'barter', '360', 'disney', 'dickie'])
def test_other_customer_exports_do_not_call_shushupapa_price_repair(monkeypatch, customer):
    def unexpected(*args, **kwargs):
        raise AssertionError('Shushupapa price repair must remain customer-specific')
    monkeypatch.setattr(mapping, '_repair_shushupapa_summary_amounts', unexpected)
    book = workbook()
    row = order(_huaxing_customer=customer)
    mapping.export_rows(book, [row])
    assert find_row(book['ITEM表'], row['reference_no']) >= 4
