"""Synthetic identifiers retain the real format boundaries without business data."""
from __future__ import annotations

from datetime import date
from io import BytesIO
from types import SimpleNamespace

import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.customer_order_360_shipping_item import (
    COMPLETED, ISSUE_CODE, ShippingItemError, backfill_shipping_items,
    derive_shipping_item, ensure_shipping_item_column, prepare_shipping_item,
)


@pytest.mark.parametrize("value, expected", [
    ("76543220069", "7654322"), (76543220069, "7654322"),
    ("8765432109.30", "8765432109"), (8765432109.3, "8765432109"),
    ("7654328", "7654328"), ("7654321098", "7654321098"),
    ("00123456789", "0012345"), ("0012345678.00", "0012345678"),
    ("0012345", "0012345"), ("0012345678", "0012345678"),
    ("76540050012a", ""), ("12345678", ""), ("123456789012", ""),
    ("1234567.1", ""), ("8765432109.", ""), ("8.765432109e9", ""),
    ("=8765432109", ""), ("１２３４５６７", ""), (True, ""), (None, ""),
])
def test_derive_only_confirmed_shapes_and_preserves_string_leading_zeroes(value, expected):
    assert derive_shipping_item(value) == expected


def _history(product, shipping, reference="RL-100-1", row=4):
    return SimpleNamespace(product_no=product, reference_no=reference, row=row,
                           source_sheet="ITEM表", extras={"shipping_item": shipping})


def _row(product="76543220069", reference="RL-100-1"):
    return {"id": "line-1", "reference_no": reference, "product_no": product, "issues": []}


def test_same_release_and_full_product_history_precedes_other_orders_and_format_rule():
    row = _row()
    prepare_shipping_item(row, [
        _history(row["product_no"], "7654003、7654004"),
        _history(row["product_no"], "7654322", "RL-100-2", 5),
    ])
    assert row["shipping_item"] == "7654003、7654004"
    assert "第4行" in row["lineage"]["shipping_item"]
    assert "第5行" not in row["lineage"]["shipping_item"]
    issue, = row["issues"]
    assert issue["can_edit"] and issue["can_skip"] and issue["severity"] == "warning"
    assert issue["skip_key"] == f"line-1|{ISSUE_CODE}|shipping_item"


def test_unique_full_product_history_may_override_rule_but_nearby_full_items_do_not():
    row = _row()
    prepare_shipping_item(row, [_history(row["product_no"], "历史组合货号", "RL-200-1")])
    assert row["shipping_item"] == "历史组合货号"
    prepare_shipping_item(row, [_history("76543220068", "其他货号"), _history("7654322", "短码另值")])
    assert row["shipping_item"] == "7654322"


@pytest.mark.parametrize("same_release", [True, False])
def test_conflicting_history_stays_blank_without_falling_back_to_rule(same_release):
    row = _row()
    prepare_shipping_item(row, [
        _history(row["product_no"], "7654003", "RL-100-1" if same_release else "RL-200-1"),
        _history(row["product_no"], "7654004", "RL-100-1" if same_release else "RL-200-2", 5),
    ])
    assert row["shipping_item"] == "" and row["status"] == "warning"
    assert "不一致" in row["lineage"]["shipping_item"]
    assert row["issues"][0]["can_edit"]


def test_prepare_recomputes_product_correction_and_replaces_only_its_own_issue():
    row = _row()
    row["issues"] = [{"severity": "blocked", "code": "quantity_conflict"}]
    prepare_shipping_item(row, [])
    assert row["shipping_item"] == "7654322"
    row["product_no"] = "UNKNOWN"
    prepare_shipping_item(row, [])
    assert row["shipping_item"] == "" and row["status"] == "blocked"
    assert len(row["issues"]) == 2
    assert row["issues"][0]["code"] == "quantity_conflict"
    assert "留空核对" in row["lineage"]["shipping_item"]


def _workbook():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "ITEM表"
    sheet["H3"] = "产品编号"
    sheet["H3"].font = Font(name="微软雅黑", bold=True)
    sheet["AL3"] = "备注"
    sheet["A10"] = "取消单"
    for row, product in ((4, "00123456789"), (5, "8765432109.30"), (6, "UNKNOWN")):
        sheet.cell(row, 5, f"RL-100-{row}")
        sheet.cell(row, 8, product)
    return workbook


def _snapshot(workbook):
    return {sheet.title: {(r, c): (cell.value, cell.data_type, cell.style_id, cell.comment.text if cell.comment else None)
                          for (r, c), cell in sheet._cells.items()}
            for sheet in workbook}


def test_free_am_and_existing_unique_reordered_header_are_used_without_shifting_columns():
    workbook = _workbook()
    sheet = workbook["ITEM表"]
    sheet["AM4"].font = Font(bold=True)  # Empty style is not business content.
    sheet["AN3"] = "版本"
    sheet["AN4"] = "两件装"
    assert ensure_shipping_item_column(workbook) == 39
    assert sheet["AM3"].value == "出货货号"
    assert sheet["AN4"].value == "两件装" and sheet["AL3"].value == "备注"
    assert sheet["AM3"].font.name == "微软雅黑" and sheet["AM3"].font.bold
    assert ensure_shipping_item_column(workbook) == 39
    sheet["AM3"] = None
    sheet["AQ3"] = " 出货\n货号 "
    assert ensure_shipping_item_column(workbook) == 43


@pytest.mark.parametrize("occupancy", ["value", "formula", "comment", "merge", "validation"])
def test_occupied_am_is_never_reused_or_moved(occupancy):
    workbook = _workbook()
    sheet = workbook["ITEM表"]
    if occupancy == "value":
        sheet["AM3"] = "版本"
        sheet["AM4"] = "人工版本"
    elif occupancy == "formula":
        sheet["AM50"] = '=IF(1=1,"","")'
    elif occupancy == "comment":
        sheet["AM20"].comment = Comment("人工保留", "业务")
    elif occupancy == "merge":
        sheet.merge_cells("AL50:AN50")
    else:
        rule = DataValidation(type="list", formula1='"甲,乙"')
        rule.add("AM4:AM100")
        sheet.add_data_validation(rule)
    before = _snapshot(workbook)["ITEM表"]
    column = ensure_shipping_item_column(workbook)
    assert column >= 40
    for key, value in before.items():
        assert _snapshot(workbook)["ITEM表"][key] == value
    assert sheet.cell(3, column).value == "出货货号"
    assert ensure_shipping_item_column(workbook) == column


def test_duplicate_shipping_headers_fail_without_changing_workbook():
    workbook = _workbook()
    sheet = workbook["ITEM表"]
    sheet["AM3"] = "出货货号"
    sheet["AN3"] = "出货货号"
    before = _snapshot(workbook)
    with pytest.raises(ShippingItemError, match="多个"):
        ensure_shipping_item_column(workbook)
    assert _snapshot(workbook) == before


def test_backfill_pending_and_explicit_completed_only_preserving_existing_values_formulas_and_identity():
    workbook = _workbook()
    sheet = workbook["ITEM表"]
    sheet["AM3"] = "出货货号"
    for row, product in ((7, "76543220069"), (8, "76543230069"), (11, "76543240069"),
                         (14, "76543250069"), (17, "76543260069"), (20, "76543270069")):
        sheet.cell(row, 5, f"RL-100-{row}")
        sheet.cell(row, 8, product)
    sheet["AM7"] = "人工出货货号"
    sheet["AM8"] = '=IF(H8="","",LEFT(H8,7))'
    sheet["A13"] = "已走货订单"
    sheet["A16"] = COMPLETED
    sheet["A19"] = "已完成"  # Ordinary legacy completion is not the explicit date group.
    identities = {r: (sheet.cell(r, 5).value, sheet.cell(r, 8).value) for r in range(4, 21)}
    report = backfill_shipping_items(workbook)
    assert report["filled"] == 3 and report["preserved"] == 2
    assert sheet["AM4"].value == "0012345"
    assert sheet["AM5"].value == "8765432109"
    assert sheet["AM17"].value == "7654326"
    assert sheet["AM7"].value == "人工出货货号" and sheet["AM8"].data_type == "f"
    for address in ("AM6", "AM11", "AM14", "AM20"):
        assert sheet[address].value is None
    assert report["unresolved"] == [{"sheet": "ITEM表", "row": 6, "cell": "AM6",
                                      "reference_no": "RL-100-6", "product_no": "UNKNOWN", "reason": "unknown"}]
    assert "ITEM表!AM6" in report["warnings"][0]
    assert identities == {r: (sheet.cell(r, 5).value, sheet.cell(r, 8).value) for r in range(4, 21)}
    assert sheet["AM4"].data_type == "s" and sheet["AM4"].number_format == "@"


def test_backfill_uses_unique_whole_product_history_and_leaves_conflict_blank():
    workbook = _workbook()
    sheet = workbook["ITEM表"]
    sheet["AM3"] = "出货货号"
    for row, product, shipping in ((11, "00123456789", "特例组合"), (12, "8765432109.30", "A"),
                                  (13, "8765432109.30", "B")):
        sheet.cell(row, 5, f"RL-200-{row}")
        sheet.cell(row, 8, product)
        sheet.cell(row, 39, shipping)
    report = backfill_shipping_items(workbook)
    assert sheet["AM4"].value == "特例组合" and sheet["AM5"].value is None
    assert next(item for item in report["unresolved"] if item["row"] == 5)["reason"] == "conflict"
    assert sheet["AM11"].value == "特例组合" and sheet["AM12"].value == "A" and sheet["AM13"].value == "B"


def test_shipping_text_is_never_executed_and_preserves_leading_zero_after_save_reload():
    workbook = _workbook()
    sheet = workbook["ITEM表"]
    sheet["AM3"] = "出货货号"
    sheet["E11"] = "RL-200-1"
    sheet["H11"] = "8765432109.30"
    sheet["AM11"] = "=手工货号"
    sheet["AM11"].data_type = "s"
    backfill_shipping_items(workbook)
    stream = BytesIO()
    workbook.save(stream)
    restored = openpyxl.load_workbook(BytesIO(stream.getvalue()))
    assert restored["ITEM表"]["AM4"].value == "0012345"
    assert restored["ITEM表"]["AM5"].value == "=手工货号"
    assert restored["ITEM表"]["AM5"].data_type == "s"
    before = _snapshot(restored)
    report = backfill_shipping_items(restored)
    assert report["filled"] == 0 and _snapshot(restored) == before


def test_target_merged_cell_is_rejected_and_example_or_repeated_header_is_ignored():
    workbook = _workbook()
    sheet = workbook["ITEM表"]
    sheet["AM3"] = "出货货号"
    sheet["AL4"] = "模拟数据"
    sheet["E7"] = "SO / Reference"
    sheet["H7"] = "产品编号"
    sheet.merge_cells("AM5:AN5")
    with pytest.raises(ShippingItemError, match="合并"):
        backfill_shipping_items(workbook)
    assert sheet["AM4"].value is None
    sheet.unmerge_cells("AM5:AN5")
    backfill_shipping_items(workbook)
    assert sheet["AM4"].value is None and sheet["AM7"].value is None


@pytest.mark.parametrize("factory, customer", [("huaxing", "360"), ("huakang-a", "green-toys"), ("huakang-d", "360")])
def test_other_customers_and_factories_are_unchanged(factory, customer):
    workbook = _workbook()
    before = _snapshot(workbook)
    assert ensure_shipping_item_column(workbook, factory, customer) is None
    assert backfill_shipping_items(workbook, factory, customer)["applied"] is False
    assert _snapshot(workbook) == before


def test_standalone_organize_twice_keeps_filled_values_and_full_product_identity():
    from app.services import customer_order_unified as unified
    from app.services.customer_order_360_workbook import organize_workbook

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name in ("ITEM表", "接单表", "正单评审表"):
        sheet = workbook.create_sheet(name)
        headers = unified.ITEM_HEADERS if name == "ITEM表" else unified.ORDER_HEADERS
        for column, header in enumerate(headers, 1):
            sheet.cell(3, column, header)
        reference, product, inspection = (5, 8, 25) if name == "ITEM表" else (4, 7, 18)
        for row, item, when in ((4, "00123456789", date(2000, 1, 1)),
                                (5, "76543220069", date(2099, 1, 1))):
            sheet.cell(row, reference, f"RL-100-{row}")
            sheet.cell(row, product, item)
            sheet.cell(row, inspection, when)
            sheet.cell(row, 2, "备料单")
        sheet["A10"] = "取消单"
        sheet["A13"] = "已走货订单"
    workbook["ITEM表"]["AM3"] = "出货货号"
    workbook["ITEM表"]["AM5"] = "手工例外"
    source = BytesIO()
    workbook.save(source)
    output, _, report, _ = organize_workbook("source.xlsx", source.getvalue())
    result = openpyxl.load_workbook(BytesIO(output))
    sheet = result["ITEM表"]
    values = {cell.value: (sheet.cell(cell.row, 8).value, sheet.cell(cell.row, 39).value)
              for cell in list(sheet._cells.values()) if cell.column == 5 and str(cell.value).startswith("RL-")}
    assert values == {"RL-100-4": ("00123456789", "0012345"), "RL-100-5": ("76543220069", "手工例外")}
    assert report["shipping_items"]["filled"] == 1
    before = _snapshot(result)
    again, _, report, _ = organize_workbook("organized.xlsx", output)
    restored = openpyxl.load_workbook(BytesIO(again))
    assert _snapshot(restored) == before and report["shipping_items"]["filled"] == 0
