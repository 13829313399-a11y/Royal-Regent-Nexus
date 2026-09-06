"""Excel updates preserve server facts and never turn allocations into shots."""

import io
import json
import os
from datetime import timedelta
from hashlib import sha256
from pathlib import Path

import pytest
from app.models.injection_scheduling import (
    Demand,
    HistoricalOutput,
    ImportBatch,
    ImportRow,
    Run,
    RunDemand,
    ShiftReport,
)
from app.services.injection_scheduling import planning
from app.services.injection_scheduling.calculations import now
from app.services.injection_scheduling.import_plan import parse_plan
from app.services.injection_scheduling.reports import refresh_quantities
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from test_injection_v3_api import BASE, ok, seed_job, write
from test_injection_v3_api import (
    env as env,  # noqa: PLC0414 -- pytest fixture re-export
)
from test_injection_v3_import_export import upload


def current_shift():
    current = now()
    return (
        (current - timedelta(days=1)).date() if current.hour < 8 else current.date()
    ).isoformat(), "DAY" if 8 <= current.hour < 20 else "NIGHT"


def export(env, *, start=None, days=14, formulas=False):
    response = env[0].post(
        BASE + "/exports/plan",
        json={
            "factory_id": "huaxing",
            "window_start": start or current_shift()[0],
            "window_days": days,
            "editable_formulas": formulas,
        },
    )
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content), data_only=False)
    sheet = book["标准交换表"]
    headers = {sheet.cell(3, col).value: col for col in range(1, sheet.max_column + 1)}
    return book, sheet, headers


def apply_book(env, book, sheet_name="标准交换表"):
    buffer = io.BytesIO()
    book.save(buffer)
    preview = ok(upload(env, buffer.getvalue(), sheet_name))
    return ok(write(env, f"/imports/{preview['batch_id']}/apply"))


def running(env):
    _, _, demand = seed_job(env)
    run = ok(write(env, "/schedule/auto", scope={"mode": "UNSCHEDULED"}))[
        "changed_runs"
    ][0]
    ok(write(env, f"/runs/{run['id']}/start"))
    return demand, run


@pytest.mark.parametrize("days", [14, 31])
def test_new_single_run_shift_then_stale_cell_conflict_keeps_unrelated_edit(env, days):
    demand, run = running(env)
    day, shift = current_shift()
    key = f"{day}|{shift}"
    book, sheet, headers = export(env, days=days)
    baseline = json.loads(sheet.cell(4, headers["roundtrip_baseline"]).value)
    assert baseline["report_targets"][key]["revision"] == 0
    sheet.cell(4, headers[key], 200)
    result = apply_book(env, book)
    assert not result["summary"]["conflicts"]
    assert result["summary"]["report_changes"][0]["physical_delta"] == 200
    book, sheet, headers = export(env, days=days)
    with Session(env[1]) as db:
        report = db.scalar(select(ShiftReport))
        report_revision = report.revision
    ok(
        write(
            env,
            "/shift-reports/current",
            method="put",
            run_id=run["id"],
            production_date=day,
            shift_code=shift,
            physical_shots=230,
            report_revision=report_revision,
        )
    )
    sheet.cell(4, headers[key], 250)
    sheet.cell(4, headers["order_note"], "仅新增备注")
    result = apply_book(env, book)
    assert [c["field"] for c in result["summary"]["conflicts"]] == [key]
    assert result["summary"]["conflicts"][0]["code"] == "REPORT_CONFLICT"
    persisted = (
        env[0]
        .get(BASE + f"/demands/{demand['id']}", params={"factory_id": "huaxing"})
        .json()["demand"]
    )
    assert (
        persisted["completed_shots"] == 230 and persisted["order_note"] == "仅新增备注"
    )


def test_unedited_old_file_preserves_report_created_after_export(env):
    demand, run = running(env)
    book, _, _ = export(env)
    day, shift = current_shift()
    ok(
        write(
            env,
            "/shift-reports/current",
            method="put",
            run_id=run["id"],
            production_date=day,
            shift_code=shift,
            physical_shots=150,
            report_revision=0,
        )
    )
    result = apply_book(env, book)
    assert (
        result["summary"]["unchanged_count"] == 1 and not result["summary"]["conflicts"]
    )
    with Session(env[1]) as db:
        assert db.get(Demand, demand["id"]).completed_shots == 150


def test_group_allocations_reject_even_forged_physical_target(env):
    demand, run = running(env)
    second = ok(
        write(
            env,
            "/demands",
            {"mold_code": demand["mold_code"], "planned_shots": 200, "order_no": "002"},
        )
    )["demand"]
    with Session(env[1]) as db:
        db.add(
            RunDemand(
                run_id=run["id"],
                factory_id="huaxing",
                demand_id=second["id"],
                sequence=2,
                allocation_mode="SEQUENTIAL_SHOTS",
                planned_quantity=200,
            )
        )
        db.commit()
    day, shift = current_shift()
    ok(
        write(
            env,
            "/shift-reports/current",
            method="put",
            run_id=run["id"],
            production_date=day,
            shift_code=shift,
            physical_shots=100,
            report_revision=0,
        )
    )
    book, sheet, headers = export(env)
    key = f"{day}|{shift}"
    row = next(
        i
        for i in range(4, sheet.max_row + 1)
        if sheet.cell(i, headers["id"]).value == demand["id"]
    )
    baseline = json.loads(sheet.cell(row, headers["roundtrip_baseline"]).value)
    assert key not in baseline["report_targets"]
    report = baseline["reports"][key][0]
    baseline["report_targets"][key] = {
        "report_key": report["report_key"],
        "revision": report["revision"],
    }
    sheet.cell(row, headers["roundtrip_baseline"], json.dumps(baseline))
    sheet.cell(row, headers[key], 150)
    result = apply_book(env, book)
    assert result["summary"]["conflicts"][0]["code"] == "INVALID_SHIFT_VALUE"
    with Session(env[1]) as db:
        assert db.scalar(select(ShiftReport)).physical_shots == 100


def test_formula_without_cache_is_not_zero_and_adjustment_in_formula(env):
    _, _, demand = seed_job(env)
    ok(write(env, f"/demands/{demand['id']}", {"adjustment_shots": 50}, method="patch"))
    book, sheet, headers = export(env, formulas=True)
    assert (
        f"{sheet.cell(3, headers['adjustment_shots']).column_letter}4"
        in sheet["N4"].value
    )
    sheet.cell(4, headers["planned_shots"], "=1000+10")
    result = apply_book(env, book)
    assert [c["code"] for c in result["summary"]["conflicts"]] == [
        "FORMULA_CACHE_MISSING"
    ]
    with Session(env[1]) as db:
        assert db.get(Demand, demand["id"]).planned_shots == 1000


def test_historical_window_does_not_move_later_reports_into_opening(env):
    _, _, demand = seed_job(env)
    start = now().date() - timedelta(days=8)
    with Session(env[1]) as db:
        batch = ImportBatch(
            factory_id="huaxing",
            file_name="test.xlsx",
            file_sha256="a" * 64,
            sheet_name="计划表",
            parser_version="fixture",
            status="APPLIED",
        )
        db.add(batch)
        db.flush()
        row = ImportRow(
            batch_id=batch.id, source_row=1, row_role="DEMAND", demand_id=demand["id"]
        )
        db.add(row)
        db.flush()
        db.get(Demand, demand["id"]).opening_shots = 10
        for offset, quantity in [(-1, 20), (0, 30), (4, 40)]:
            db.add(
                HistoricalOutput(
                    demand_id=demand["id"],
                    import_row_id=row.id,
                    production_date=start + timedelta(days=offset),
                    shift_code="DAY",
                    source_key="legacy",
                    quantity=quantity,
                    counts_toward_demand=True,
                )
            )
        db.flush()
        refresh_quantities(db, "huaxing")
        db.commit()
    _, sheet, headers = export(env, start=start.isoformat(), days=2)
    assert sheet.cell(4, headers["window_opening_shots"]).value == 30
    assert sheet.cell(4, headers["completed_shots"]).value == 60
    assert sheet.cell(4, headers["current_completed_shots"]).value == 100
    assert sheet.cell(4, headers["remaining_shots"]).value == 940


def legacy_book(
    *, planned=1000, completed=100, history=50, note="旧备注", duplicate=False
):
    book = Workbook()
    sheet = book.active
    sheet.title = "计划表"
    sheet.cell(1, 1, "日计划")
    day = now().date() - timedelta(days=2)
    sheet.cell(2, 13, day)
    headers = [
        "模号",
        "单号",
        "货号",
        "计划啤数",
        "已啤数",
        "欠数",
        "计划日目标",
        "颜色",
        "色粉",
        "用料",
        "工艺备注",
        "模具机安 A",
        "白班",
        "夜班",
    ]
    for c, key in enumerate(headers, 1):
        sheet.cell(3, c, key)
    for r in range(4, 6 if duplicate else 5):
        values = [
            "LEG-01",
            "ORD01",
            "ITEM01",
            planned if r == 4 else 500,
            completed,
            planned - completed,
            2400,
            "白",
            "W1",
            "PP",
            note,
            12,
            history,
            None,
        ]
        for c, value in enumerate(values, 1):
            sheet.cell(r, c, value)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def apply_legacy(env, content):
    preview = ok(upload(env, content))
    result = ok(write(env, f"/imports/{preview['batch_id']}/apply"))
    return preview, result


def test_legacy_daily_revision_updates_split_rows_without_duplicates(env):
    _, first = apply_legacy(env, legacy_book(duplicate=True))
    assert first["summary"]["applied_count"] == 2
    preview, result = apply_legacy(
        env, legacy_book(planned=1100, completed=120, history=70, duplicate=True)
    )
    assert preview["summary"]["updates"] == 2 and preview["summary"]["creates"] == 0
    assert not result["summary"]["conflicts"]
    with Session(env[1]) as db:
        rows = list(db.scalars(select(Demand).order_by(Demand.source_row)))
        assert len(rows) == 2
        assert [d.planned_shots for d in rows] == [1100, 500]
        assert [d.completed_shots for d in rows] == [120, 120]
        assert [d.opening_shots for d in rows] == [50, 50]
        assert len(list(db.scalars(select(HistoricalOutput)))) == 2


def test_legacy_conflict_only_field_server_changed_other_fields_apply(env):
    _, first = apply_legacy(env, legacy_book())
    demand_id = first["demand_ids"][0]
    ok(
        write(
            env,
            f"/demands/{demand_id}",
            {"production_note": "服务器备注"},
            method="patch",
        )
    )
    _, result = apply_legacy(env, legacy_book(planned=1100, note="文件新备注"))
    assert [(c["code"], c["field"]) for c in result["summary"]["conflicts"]] == [
        ("FIELD_CONFLICT", "production_note")
    ]
    with Session(env[1]) as db:
        demand = db.get(Demand, demand_id)
        assert demand.planned_shots == 1100 and demand.production_note == "服务器备注"


def test_conflicting_duplicate_physical_report_cells_both_stay_unapplied(env):
    demand, _ = running(env)
    book, sheet, headers = export(env)
    key = "|".join(current_shift())
    for col in range(1, sheet.max_column + 1):
        sheet.cell(5, col, sheet.cell(4, col).value)
    sheet.cell(4, headers[key], 200)
    sheet.cell(5, headers[key], 250)
    result = apply_book(env, book)
    assert [c["code"] for c in result["summary"]["conflicts"]] == [
        "DUPLICATE_REPORT_CONFLICT"
    ] * 2
    assert result["summary"]["report_changes"] == []
    with Session(env[1]) as db:
        assert db.get(Demand, demand["id"]).completed_shots == 0


def test_long_window_metadata_survives_excel_cell_text_limit(env):
    _, run = running(env)
    start = now() - timedelta(days=300)
    with Session(env[1]) as db:
        db.get(Run, run["id"]).actual_start_at = start
        db.commit()
    book, sheet, headers = export(env, start=start.date().isoformat(), days=301)
    assert sheet.cell(4, headers["roundtrip_baseline_2"]).value
    output = io.BytesIO()
    book.save(output)
    parsed = parse_plan(output.getvalue(), "标准交换表")
    assert not parsed["demands"][0]["issues"]
    assert len(parsed["demands"][0]["baseline"]["report_targets"]) >= 599
    assert parsed["read_worksheets"] == ["xl/worksheets/sheet2.xml"]


def test_exchange_preview_prefers_stable_id_when_order_identity_is_edited(env):
    _, _, demand = seed_job(env)
    book, sheet, headers = export(env)
    sheet.cell(4, headers["order_no"], "UPDATED-ORDER")
    output = io.BytesIO()
    book.save(output)
    preview = ok(upload(env, output.getvalue(), "标准交换表"))
    assert preview["summary"]["updates"] == 1 and preview["summary"]["creates"] == 0
    applied = ok(write(env, f"/imports/{preview['batch_id']}/apply"))
    assert (
        applied["summary"]["applied_count"] == 1 and not applied["summary"]["conflicts"]
    )
    with Session(env[1]) as db:
        assert db.get(Demand, demand["id"]).order_no == "UPDATED-ORDER"


def test_co_output_shift_exports_good_units_without_adding_to_shot_m(env):
    _, _, demand = seed_job(env)
    ok(
        write(
            env,
            f"/demands/{demand['id']}",
            {
                "allocation_mode": "CO_OUTPUT_UNITS",
                "required_units": 200,
                "effective_outputs_per_shot": 2,
                "output_configuration": {"M-01": 2},
            },
            method="patch",
        )
    )
    run = ok(write(env, "/schedule/auto", scope={"mode": "UNSCHEDULED"}))[
        "changed_runs"
    ][0]
    ok(write(env, f"/runs/{run['id']}/start"))
    day, shift = current_shift()
    ok(
        write(
            env,
            "/shift-reports/current",
            method="put",
            run_id=run["id"],
            production_date=day,
            shift_code=shift,
            physical_shots=10,
            report_revision=0,
        )
    )
    book, sheet, headers = export(env)
    key = f"{day}|{shift}"
    assert sheet.cell(4, headers[key]).value == 20
    assert sheet.cell(4, headers["shift_quantity_unit"]).value == "UNIT"
    assert sheet.cell(4, headers["completed_shots"]).value == 0
    assert sheet.cell(4, headers["good_units"]).value == 20
    assert sheet.cell(4, headers["current_good_units"]).value == 20
    sheet.cell(4, headers[key], 22)
    result = apply_book(env, book)
    assert (
        result["summary"]["conflicts"][0]["code"] == "SHIFT_NOT_UNIQUE_PHYSICAL_REPORT"
    )


def test_real_source_daily_hash_change_and_14_31_day_roundtrips(env):
    path = os.getenv("INJECTION_V3_WORKBOOK")
    if not path:
        pytest.skip("Set INJECTION_V3_WORKBOOK for private source regression")
    source = Path(path)
    content = source.read_bytes()
    digest = sha256(content).hexdigest()
    _, first = apply_legacy(env, content)
    assert first["summary"]["applied_count"] == 280
    originally_scheduled = stored_run_times(env)
    # A new transport hash with exactly the same selected-sheet cells exercises
    # source identity; the original archive remains byte-for-byte untouched.
    preview, second = apply_legacy(env, content + b"\nV3 read-only regression copy\n")
    assert preview["summary"]["creates"] == 0 and preview["summary"]["updates"] == 280
    assert (
        second["summary"]["applied_count"] == 0
        and second["summary"]["unchanged_count"] == 280
    )
    assert not second["summary"]["conflicts"]
    assert second["changed_runs"] == [] and second["recalculate_required"] is False
    assert stored_run_times(env) == originally_scheduled
    for days in (14, 31):
        book, _, _ = export(env, start="2026-09-01", days=days)
        for sheet_name in ("计划表", "标准交换表"):
            result = apply_book(env, book, sheet_name)
            assert result["summary"]["unchanged_count"] == 280
            assert not result["summary"]["conflicts"]
            assert (
                result["changed_runs"] == [] and result["recalculate_required"] is False
            )
            assert stored_run_times(env) == originally_scheduled
    with Session(env[1]) as db:
        assert len(list(db.scalars(select(Demand)))) == 280
        assert len(list(db.scalars(select(HistoricalOutput)))) == 1449
    assert sha256(source.read_bytes()).hexdigest() == digest


def stored_run_times(env):
    with Session(env[1]) as db:
        return [
            (
                r.id,
                r.machine_id,
                r.mold_asset_id,
                r.status,
                r.setup_start_at,
                r.planned_start_at,
                r.planned_end_at,
                r.sequence,
                r.revision,
            )
            for r in db.scalars(select(Run).order_by(Run.id))
        ]


@pytest.mark.parametrize("sheet_name", ["计划表", "标准交换表"])
def test_unchanged_scheduled_export_import_does_not_move_runs(
    env, monkeypatch, sheet_name
):
    _, _, demand = seed_job(env)
    ok(
        write(
            env,
            "/demands",
            {
                "mold_code": demand["mold_code"],
                "order_no": "FOLLOWING",
                "planned_shots": 500,
                "material_raw": "PP",
            },
        )
    )
    ok(write(env, "/schedule/auto", scope={"mode": "UNSCHEDULED"}))
    before = stored_run_times(env)
    assert len(before) == 2
    book, _, _ = export(env)
    future = now() + timedelta(hours=2)
    monkeypatch.setattr(planning, "now", lambda: future)
    result = apply_book(env, book, sheet_name)
    assert result["summary"]["applied_count"] == 0
    assert result["summary"]["unchanged_count"] == 2
    assert result["summary"]["conflicts"] == []
    assert result["recalculate_required"] is False
    assert result["changed_runs"] == []
    assert stored_run_times(env) == before


def test_all_conflicting_import_keeps_schedule_but_applied_quantity_recalculates(
    env, monkeypatch
):
    _, _, demand = seed_job(env)
    ok(write(env, "/schedule/auto", scope={"mode": "UNSCHEDULED"}))
    book, sheet, headers = export(env)
    ok(write(env, f"/demands/{demand['id']}", {"planned_shots": 1500}, method="patch"))
    before = stored_run_times(env)
    sheet.cell(4, headers["planned_shots"], 2000)
    future = now() + timedelta(hours=2)
    monkeypatch.setattr(planning, "now", lambda: future)
    result = apply_book(env, book)
    assert result["summary"]["conflicts"][0]["code"] == "FIELD_CONFLICT"
    assert result["summary"]["applied_count"] == 0
    assert result["changed_runs"] == []
    assert result["recalculate_required"] is False
    assert stored_run_times(env) == before
    # A partial conflict with a separately accepted quantity adjustment still
    # changes the authoritative remaining amount and recalculates the schedule.
    sheet.cell(4, headers["adjustment_shots"], 50)
    partial = apply_book(env, book)
    assert partial["summary"]["applied_count"] == 1
    assert partial["summary"]["conflicts"][0]["code"] == "FIELD_CONFLICT"
    assert partial["recalculate_required"] is True and partial["changed_runs"]
    assert stored_run_times(env) != before
    with Session(env[1]) as db:
        assert db.get(Demand, demand["id"]).remaining_shots == 1550


def test_unchanged_legacy_hash_keeps_runs_but_new_machine_requests_recalculation(
    env, monkeypatch
):
    content = legacy_book()
    _, imported = apply_legacy(env, content)
    ok(
        write(
            env,
            f"/demands/{imported['demand_ids'][0]}",
            {"required_machine_a": 12},
            method="patch",
        )
    )
    ok(
        write(
            env,
            "/machines",
            {"code": "01", "machine_a": 12, "machine_family": "HORIZONTAL"},
        )
    )
    scheduled = ok(write(env, "/schedule/auto", scope={"mode": "UNSCHEDULED"}))
    before = stored_run_times(env)
    assert before, scheduled
    future = now() + timedelta(hours=2)
    monkeypatch.setattr(planning, "now", lambda: future)
    _, unchanged = apply_legacy(env, content + b"\nnew transport version\n")
    assert unchanged["summary"]["applied_count"] == 0
    assert unchanged["summary"]["created_machine_ids"] == []
    assert (
        unchanged["recalculate_required"] is False and unchanged["changed_runs"] == []
    )
    assert stored_run_times(env) == before
    # A machine-only import is still a real resource change.
    with Session(env[1]) as db:
        batch = ImportBatch(
            factory_id="huaxing",
            file_name="machine-only-plan.xlsx",
            file_sha256="b" * 64,
            sheet_name="计划表",
            parser_version="injection-v3-plan-1",
            status="PREVIEW",
            evidence={
                "machines": [
                    {"code": "02", "machine_a": 12, "machine_family": "HORIZONTAL"}
                ],
                "structure_rows": [],
                "demands": [],
            },
            summary={},
        )
        db.add(batch)
        db.flush()
        batch_id = batch.id
        db.commit()
    changed = ok(write(env, f"/imports/{batch_id}/apply"))
    assert changed["summary"]["applied_count"] == 0
    assert len(changed["summary"]["created_machine_ids"]) == 1
    assert changed["recalculate_required"] is True and changed["changed_runs"]
