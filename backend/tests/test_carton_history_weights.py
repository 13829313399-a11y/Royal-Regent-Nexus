"""Packed-goods weights stay on their declared paper through history import."""
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.services.carton_procurement_history_import import _parse_rows
from test_carton_history_wide import MINIMAL_HEADERS, row, workbook


WEIGHT_HEADERS = ["外箱每箱净重kg", "外箱每箱毛重kg"]
ROOT = Path(__file__).resolve().parents[2]


def test_main_weights_apply_only_to_outer_carton_and_old_templates_remain_unknown():
    data = row(内箱需求数量=200, 卡纸需求数量=100, 外箱每箱净重kg="8.125", 外箱每箱毛重kg="9.25")
    parsed, _ = _parse_rows("history.xlsx", workbook([data], MINIMAL_HEADERS + WEIGHT_HEADERS))
    papers = {entry["line"].packaging_type: entry["line"] for entry in parsed}
    assert (papers["外箱"].net_weight_kg, papers["外箱"].gross_weight_kg) == (Decimal("8.125"), Decimal("9.25"))
    assert all(papers[key].net_weight_kg is None and papers[key].gross_weight_kg is None for key in ("内箱", "卡纸"))
    old, _ = _parse_rows("old.xlsx", workbook([data], MINIMAL_HEADERS))
    assert all(entry["line"].net_weight_kg is None and entry["line"].gross_weight_kg is None for entry in old)


@pytest.mark.parametrize("net_header,gross_header", [("净重", "毛重"), ("每箱净重（kg）", "每箱毛重（kg）")])
def test_unqualified_main_weights_mean_outer_carton_and_preserve_zero(net_header, gross_header):
    parsed, _ = _parse_rows("history.xlsx", workbook([row(**{net_header: 0, gross_header: 0})], MINIMAL_HEADERS + [net_header, gross_header]))
    assert parsed[0]["line"].net_weight_kg == parsed[0]["line"].gross_weight_kg == 0


@pytest.mark.parametrize("net,gross", [("-1", "2"), ("9", "8"), ("heavy", "9"), ("NaN", "9"),
                                       ("1.00001", "2"), ("100000000", "100000001")])
def test_bad_weights_block_instead_of_becoming_unknown_or_rounded(net, gross):
    with pytest.raises(HTTPException) as error:
        _parse_rows("history.xlsx", workbook([row(外箱每箱净重kg=net, 外箱每箱毛重kg=gross)], MINIMAL_HEADERS + WEIGHT_HEADERS))
    assert error.value.status_code == 422 and "重" in error.value.detail


@pytest.mark.parametrize("quantity", [0, None])
def test_outer_weight_without_corresponding_demand_is_not_silently_dropped(quantity):
    with pytest.raises(HTTPException, match="对应纸品需求数量"):
        _parse_rows("history.xlsx", workbook([row(外箱需求数量=quantity, 内箱需求数量=100, 外箱每箱净重kg=1, 外箱每箱毛重kg=2)], MINIMAL_HEADERS + WEIGHT_HEADERS))


def test_duplicate_outer_weight_headers_block_ambiguous_reading():
    with pytest.raises(HTTPException, match="出现多列"):
        _parse_rows("history.xlsx", workbook([row(净重=1, 外箱每箱净重kg=2, 外箱每箱毛重kg=3)], MINIMAL_HEADERS + ["净重"] + WEIGHT_HEADERS))


def test_explicit_inner_weights_stay_separate_from_main_outer_weights():
    parsed, _ = _parse_rows("history.xlsx", workbook([row(内箱需求数量=100, 外箱每箱净重kg=8, 外箱每箱毛重kg=9,
        内箱每箱净重kg=1, 内箱每箱毛重kg=2)], MINIMAL_HEADERS + WEIGHT_HEADERS + ["内箱每箱净重kg", "内箱每箱毛重kg"]))
    assert {entry["line"].packaging_type: (entry["line"].net_weight_kg, entry["line"].gross_weight_kg) for entry in parsed} == {
        "内箱": (1, 2), "外箱": (8, 9)}


def test_extra_paper_and_legacy_vertical_import_read_both_weights():
    from test_carton_history_identity_api import _row
    from test_carton_procurement_api import _history_workbook_bytes
    legacy = load_workbook(BytesIO(_history_workbook_bytes([_row("迪奇")])))
    sheet = legacy.active
    column = sheet.max_column + 1
    for offset, (title, value) in enumerate(zip(["每箱净重kg", "每箱毛重kg"], [1.2, 1.5])):
        sheet.cell(5, column + offset, title)
        sheet.cell(6, column + offset, value)
    stream = BytesIO(); legacy.save(stream)
    parsed, _ = _parse_rows("legacy.xlsx", stream.getvalue())
    assert (parsed[0]["line"].net_weight_kg, parsed[0]["line"].gross_weight_kg) == (Decimal("1.2"), Decimal("1.5"))
    extra = load_workbook(BytesIO(workbook([row(历史订单号="H-EXTRA")], MINIMAL_HEADERS + ["历史订单号"],
        extras=[["H-EXTRA", "内箱", 50, "B3", "10*10*10", "个"]])))
    sheet = extra["附加纸品明细"]
    sheet.cell(1, 7, "每箱净重kg"); sheet.cell(1, 8, "每箱毛重kg")
    sheet.cell(2, 7, 0.3); sheet.cell(2, 8, 0.4)
    stream = BytesIO(); extra.save(stream)
    parsed, _ = _parse_rows("extra.xlsx", stream.getvalue())
    assert parsed[0]["line"].net_weight_kg is None
    assert (parsed[1]["line"].net_weight_kg, parsed[1]["line"].gross_weight_kg) == (Decimal("0.3"), Decimal("0.4"))


def test_published_template_weights_preview_save_and_invalid_batch_are_atomic(monkeypatch):
    from test_molding_sample_api import make_client, login_as
    from test_carton_history_identity_api import _customer, _orders
    from test_carton_history_preview_api import upload
    template = load_workbook(ROOT / "public/templates/carton-history-order-import-template.xlsx")
    assert [cell.value for cell in template["历史订单导入"][2]][8:10] == WEIGHT_HEADERS
    for target, source in ((3, 6), (4, 7)):
        for column, cell in enumerate(template["填写示例"][source], 1):
            template["历史订单导入"].cell(target, column, cell.value)
    valid = BytesIO(); template.save(valid)
    template["历史订单导入"].cell(4, 10, 1)  # Second order: gross below net.
    invalid = BytesIO(); template.save(invalid)
    with make_client(monkeypatch) as client:
        login_as(client, "admin"); _customer(client, "迪奇")
        bad = upload(client, invalid.getvalue()).json()
        assert bad["errors"] and "毛重" in " ".join(bad["errors"])
        assert upload(client, invalid.getvalue(), preview=False, fingerprint=bad["source_fingerprint"]).status_code == 422
        assert _orders(client) == []
        preview = upload(client, valid.getvalue()).json()
        assert preview["errors"] == []
        assert Decimal(preview["orders"][0]["lines"][0]["net_weight_kg"]) == Decimal("8.125")
        assert Decimal(preview["orders"][0]["lines"][0]["gross_weight_kg"]) == Decimal("9.25")
        saved = upload(client, valid.getvalue(), preview=False, fingerprint=preview["source_fingerprint"])
        assert saved.status_code == 201, saved.text
        orders = _orders(client)
        assert len(orders) == 2
        for order in orders:
            assert order["net_weight_kg"] is None and order["gross_weight_kg"] is None
            for line in order["lines"]:
                if line["packaging_type"] == "外箱":
                    assert line["net_weight_kg"] is not None and Decimal(line["gross_weight_kg"]) > Decimal(line["net_weight_kg"])
                else:
                    assert line["net_weight_kg"] is None and line["gross_weight_kg"] is None
        assert client.get("/api/carton-procurement/inventory/movements", params={"factory_id": "huaxing"}).json()["total"] == 0
