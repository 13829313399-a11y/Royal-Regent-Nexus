from copy import deepcopy
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

from test_carton_order_import import isolated, make_client, setup, regular, special_row, read_body, checked, post, ok


def saved(book):
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def populated(mode, *, count=1):
    from app.services.carton_order_import_file import template, fields
    book = load_workbook(BytesIO(template(mode)))
    row = regular() if mode == 'REGULAR' else special_row()
    keys = [key for key, _, _ in fields(mode)]
    for n in range(3, count + 3):
        for col, key in enumerate(keys, 2):
            book.active.cell(n, col, row.get(key, ''))
    return book


@pytest.mark.parametrize('mode', ['REGULAR', 'SPECIAL'])
def test_template_requirements_and_demo_round_trip(mode):
    from app.services.carton_order_import_file import template, read, field_specs, DEMO_COLUMN
    empty = read(template(mode), mode)
    assert empty['rows'] == [] and empty['ignored_demo_rows'] == [2]
    specs = {f['key']: f for f in empty['fields']}
    assert all(f.get('hint') and f['requirement'] in {'required', 'optional', 'conditional'} for f in specs.values())
    assert specs['specification']['requirement'] == 'required'
    assert specs['note']['requirement'] == 'optional'
    assert specs['quantity']['requirement'] == ('conditional' if mode == 'REGULAR' else 'required')
    assert specs['group']['requirement'] == ('optional' if mode == 'REGULAR' else 'required')
    assert len(empty['mapping']) == len(field_specs(mode))
    assert not any(c['label'] == DEMO_COLUMN for c in empty['columns'])
    parsed = read(saved(populated(mode)), mode)
    assert [r['source_row'] for r in parsed['rows']] == [3]
    assert parsed['ignored_demo_rows'] == [2]
    assert parsed['rows'][0]['data']['paper_quality'] == 'A33'
    with pytest.raises(HTTPException, match='列匹配无效'):
        read(template(mode), mode, mapping={'note': 1})


@pytest.mark.parametrize('mode', ['REGULAR', 'SPECIAL'])
def test_demo_column_reordering_and_edited_demo_never_silently_discards_data(mode):
    from app.services.carton_order_import_file import read, fields, DEMO_MARKER
    original = populated(mode)
    book = Workbook(); sheet = book.active
    sheet.append(['说明']); sheet.append([])
    for row in original.active.iter_rows(values_only=True):
        sheet.append(list(reversed(row)))
    original.close()
    result = read(saved(book), mode)
    assert result['header_row'] == 3
    assert result['ignored_demo_rows'] == [4]
    assert [r['source_row'] for r in result['rows']] == [5]

    book = populated(mode)
    quality_column = [k for k, _, _ in fields(mode)].index('paper_quality') + 2
    book.active.cell(2, quality_column, '真实纸质')
    changed = saved(book)
    with pytest.raises(HTTPException, match='仍带演示标记'):
        read(changed, mode)
    book = load_workbook(BytesIO(changed)); book.active.cell(2, 1).value = None
    result = read(saved(book), mode)
    assert [r['source_row'] for r in result['rows']] == [2, 3]
    assert result['ignored_demo_rows'] == []
    assert result['rows'][0]['data']['paper_quality'] == '真实纸质'

    book = populated(mode)
    note_column = [k for k, _, _ in fields(mode)].index('note') + 2
    book.active.cell(3, note_column, DEMO_MARKER)
    result = read(saved(book), mode)
    assert len(result['rows']) == 1 and result['rows'][0]['data']['note'] == DEMO_MARKER


def test_500_real_rows_plus_demo_at_header_50_preserve_physical_source_numbers():
    from app.services.carton_order_import_file import read
    from app.schemas.carton_order_import import ImportRow
    from app.services.carton_demand_import import ImportRow as SpecialRow, KEYS
    book = populated('SPECIAL', count=500)
    book.active.insert_rows(1, 49)
    result = read(saved(book), 'SPECIAL')
    assert result['header_row'] == 50 and result['ignored_demo_rows'] == [51]
    assert len(result['rows']) == 500
    last = result['rows'][-1]
    assert last['source_row'] == 551
    assert ImportRow.model_validate(last).source_row == 551
    assert SpecialRow(source_row=551, **{k: last['data'][k] for k in KEYS}).source_row == 551
    with pytest.raises(HTTPException, match='500 行'):
        read(saved(populated('SPECIAL', count=501)), 'SPECIAL')


@pytest.mark.parametrize('mode', ['REGULAR', 'SPECIAL'])
def test_api_creates_only_actual_order_and_rejects_reinserting_demo_source(monkeypatch, mode):
    with make_client(monkeypatch) as client:
        setup(client, monkeypatch)
        content = saved(populated(mode))
        body = read_body(client, content, mode)
        assert [r['source_row'] for r in body['rows']] == [3]
        command, preview = checked(client, content, body)
        assert len(preview['groups']) == 1
        assert preview['groups'][0]['source_rows'] == [3]
        assert preview['groups'][0]['quantities'] == (['101.0000'] if mode == 'REGULAR' else ['12.125'])
        forged = deepcopy(command)
        forged['rows'].append({'source_row': 2, 'data': dict(body['rows'][0]['data'])})
        assert post(client, 'confirm', content, forged).status_code == 422
        result = ok(post(client, 'confirm', content, command))
        assert result['count'] == 1
        assert ok(post(client, 'confirm', content, command)) == result
