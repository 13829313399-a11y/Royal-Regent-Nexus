from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from io import BytesIO
from threading import Barrier
from uuid import uuid4
import importlib.util
import json
from pathlib import Path
import pytest
from test_molding_sample_api import make_client, login_as
from test_carton_procurement_api import _create_order, _order_payload, _submit_order

BASE = "/api/carton-supplier"
PASSWORD = "SupplierOnly123!"

def grant_supplier_permissions(user_id="supplier-test", permissions=(
    "carton_supplier:read", "carton_supplier:edit", "carton_supplier:approve")):
    from sqlalchemy import select
    from app.db import SessionLocal
    from app.models.auth import AuthPermission, AuthUserPermissionOverride
    with SessionLocal() as db:
        for code in permissions:
            permission = db.scalar(select(AuthPermission).where(AuthPermission.code == code))
            assert permission is not None, code
            db.add(AuthUserPermissionOverride(id=f"supplier-allow-{user_id}-{code}", user_id=user_id,
                permission_id=permission.id, effect="allow", factory_id="*", department="*",
                status="active", reason="测试供应商模块权限"))
        db.commit()

def revoke_supplier_permission(code="carton_supplier:read", user_id="supplier-test"):
    from app.db import SessionLocal
    from app.models.auth import AuthUserPermissionOverride
    with SessionLocal() as db:
        row = db.get(AuthUserPermissionOverride, f"supplier-allow-{user_id}-{code}")
        assert row is not None
        row.status = "revoked"
        db.commit()

def restore_supplier_permission(code, user_id="supplier-test"):
    from app.db import SessionLocal
    from app.models.auth import AuthUserPermissionOverride
    with SessionLocal() as db:
        row = db.get(AuthUserPermissionOverride, f"supplier-allow-{user_id}-{code}")
        assert row is not None
        row.status = "active"
        db.commit()

def setup_portal(client):
    login_as(client, "admin")
    order = _create_order(client)
    response = client.post(f"/api/carton-procurement/orders/{order['order_no']}/purchase-order-issues.xlsx",
        json={"factory_id": "huaxing", "expected_revision": order["revision"]})
    assert response.status_code == 200, response.text
    from app.db import SessionLocal
    from app.models.auth import AuthUser, AuthUserAuthorizationRevision, EmployeeProfile
    from app.services.auth import make_password_hash
    salt, digest = make_password_hash(PASSWORD)
    with SessionLocal() as db:
        db.add(AuthUser(id="supplier-test", username="supplier-test", display_name="供应商测试", password_salt=salt,
            password_hash=digest, status="active", force_password_change=0, created_at="2026-09-21", updated_at="2026-09-21"))
        db.add(AuthUserAuthorizationRevision(user_id="supplier-test", revision=1, updated_at="2026-09-21"))
        db.add(EmployeeProfile(user_id="supplier-test", primary_factory_id="huaxing",
            primary_department="pmc-warehouse", position="供应商联系人",
            confirmation_status="confirmed", created_at="2026-09-21", updated_at="2026-09-21"))
        db.commit()
    grant_supplier_permissions()
    supplier = supplier_login(client)
    assert {"carton_supplier:read", "carton_supplier:edit", "carton_supplier:approve"}.issubset(supplier["permissions"])
    return order

def supplier_login(client, *, expect_read=True):
    result = client.post("/api/auth/login", json={"username":"supplier-test", "password":PASSWORD})
    assert result.status_code == 200, result.text
    assert "carton_procurement:read" not in result.json()["permissions"]
    has_read = any(item["permission_code"] == "carton_supplier:read" and item["factory_id"] == "*"
                   and item["department"] == "*" and item["effect"] == "allow"
                   for item in result.json()["effective_access"])
    assert has_read is expect_read, [item for item in result.json()["effective_access"]
                                    if item["permission_code"] == "carton_supplier:read"]
    return result.json()

def accept_all(client):
    workspace = client.get(BASE+"/workspace", params={"factory_id":"huaxing"}).json()
    order = workspace["orders"][0]
    for line in order["lines"]:
        response = client.put(BASE+f"/papers/{line['id']}/commitment", json={"factory_id":"huaxing", "issue_id":order["issue_id"], "expected_revision":0, "promised_date":"2026-09-23"})
        assert response.status_code == 200, response.text
    return order

def ship_payload(order):
    return {"factory_id":"huaxing", "request_id":str(uuid4()), "delivery_note_no":"DN-PORTAL", "delivery_date":"2026-09-21",
        "lines":[{"order_line_id":line["id"], "issue_id":order["issue_id"], "quantity":10} for line in order["lines"]]}

def receive_payload(client, shipment):
    locations = client.get("/api/carton-procurement/inventory/locations", params={"factory_id":"huaxing"}).json()
    location = next(row["id"] for row in locations if row["bin_code"] == "A-01")
    return {"factory_id":"huaxing", "request_id":str(uuid4()), "expected_revision":shipment["revision"], "acceptance_date":"2026-09-21",
        "lines":[{"shipment_line_id":line["id"], "received_quantity":8, "unit_price":2,
            "paper_quality":line["paper_quality"], "specification":line["specification"],
            "location_allocations":[{"location_id":location,"quantity":8}], "difference_reason":"本次实际短收两件"} for line in shipment["lines"]]}


def _dongkang_delivery_file(order, *, missing_item=False, destination="华兴", source_spec=None):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "送货明细"
    sheet.append(["送货单号", "日期", "客户", "客户单号", "客户料号", "纸质", "规格", "数量"])
    for index, line in enumerate(order["lines"]):
        sheet.append(["DK-IMPORT-01", "2026-09-24", destination, order["contract_no"],
            "" if missing_item and index == 0 else order["item_no"], line["paper_quality"],
            source_spec or line["specification"], 10])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_supplier_mixed_delivery_requires_warehouse_no_order_decision(monkeypatch):
    from openpyxl import Workbook
    with make_client(monkeypatch) as client:
        setup_portal(client)
        order = accept_all(client)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "送货明细"
        sheet.append(["送货单号", "日期", "客户", "客户单号", "客户料号", "名称", "纸质", "规格", "数量"])
        paper = order["lines"][0]
        sheet.append(["DK-MIXED-01", "2026-09-24", "华兴", order["contract_no"], order["item_no"],
            paper["packaging_type"], paper["paper_quality"], paper["specification"], 5])
        sheet.append(["DK-MIXED-01", "2026-09-24", "华兴", "SAMPLE-1", "SAMPLE-BOX", "外箱", "A33+B", "18*12.5*17.25 cm", 4])
        sheet.append(["DK-MIXED-01", "2026-09-24", "华兴", "WRONG-1", "WRONG-BOX", "外箱", "A33+B", "12*10*5 cm", 3])
        buffer = BytesIO(); workbook.save(buffer)
        upload = {"file": ("东康送货明细.xlsx", buffer.getvalue())}
        preview = client.post(BASE + "/shipments/import-preview", files=upload)
        assert preview.status_code == 200, preview.text
        group = preview.json()["groups"][0]
        assert group["ready"] is True
        assert [line["status"] for line in group["rows"]] == ["READY", "AD_HOC_REVIEW", "AD_HOC_REVIEW"]
        sent = client.post(BASE + "/shipments/import-confirm", files=upload, data={
            "sha256": preview.json()["sha256"],
            "selections": json.dumps([{"factory_id": "huaxing", "delivery_note_no": "DK-MIXED-01"}])})
        assert sent.status_code == 200, sent.text
        shipment = sent.json()["shipments"][0]
        assert len(shipment["lines"]) == 3
        assert sum(line["source_type"] == "AD_HOC_REVIEW" for line in shipment["lines"]) == 2

        login_as(client, "warehouse_keeper")
        locations = client.get("/api/carton-procurement/inventory/locations", params={"factory_id": "huaxing"}).json()
        bin_id = next(location["id"] for location in locations if location["bin_code"] == "A-01")
        lines = []
        for line in shipment["lines"]:
            quantity = int(Decimal(line["quantity"]))
            data = {"shipment_line_id": line["id"], "received_quantity": quantity,
                "unit_price": "2.5", "paper_quality": line["paper_quality"],
                "specification": line["specification"], "location_allocations": [{"location_id": bin_id, "quantity": quantity}],
                "difference_reason": ""}
            if line["item_no"] == "SAMPLE-BOX":
                data.update(no_order_decision="SAMPLE", customer_code="DICKIE",
                    sample_purpose="客户打板确认", requested_by="纸箱部")
            if line["item_no"] == "WRONG-BOX":
                data.update(no_order_decision="WRONG_DELIVERY", rejected_quantity=quantity,
                    location_allocations=[], difference_reason="供应商送错纸品")
            lines.append(data)
        payload = {"factory_id": "huaxing", "request_id": str(uuid4()), "expected_revision": shipment["revision"],
            "acceptance_date": "2026-09-24", "lines": lines}
        unsafe = dict(payload, lines=[dict(line, no_order_decision="") if line["shipment_line_id"] ==
            next(item["id"] for item in shipment["lines"] if item["item_no"] == "SAMPLE-BOX") else line for line in lines])
        denied = client.post(BASE + f"/internal/shipments/{shipment['id']}/receive", json=unsafe)
        assert denied.status_code == 422, denied.text
        posted = client.post(BASE + f"/internal/shipments/{shipment['id']}/receive", json=payload)
        assert posted.status_code == 200, posted.text
        assert posted.json()["status"] == "RECEIVED"
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonReceipt, CartonReceiptLine, CartonInventoryMovement
        from app.services.carton_supplier_settlement import sources
        from sqlalchemy import select
        with SessionLocal() as db:
            receipt_lines = db.scalars(select(CartonReceiptLine).where(CartonReceiptLine.receipt_id == posted.json()["receipt_id"])).all()
            assert {line.source_type for line in receipt_lines} == {"FORMAL_ORDER", "AD_HOC"}
            assert len(receipt_lines) == 2
            assert any(line.source_type == "AD_HOC" and line.order_line_id is None and "客户打板确认" in line.feedback_note for line in receipt_lines)
            movements = db.scalars(select(CartonInventoryMovement).where(CartonInventoryMovement.source_id == posted.json()["receipt_id"])).all()
            assert len(movements) == 2
            receipt = db.get(CartonReceipt, posted.json()["receipt_id"])
            settlement_rows, _, settlement_issues, undated = sources(db, "huaxing", receipt.supplier_id, "2026-09", "CNY")
            assert not settlement_issues and not undated
            assert {row["item_no"] for row in settlement_rows} == {order["item_no"], "SAMPLE-BOX"}
            assert all(row["unit_price"] is not None for row in settlement_rows)

        # A later formal order can consume the original sample evidence without
        # posting another receipt, inventory movement, or month-end source.
        sample = next(line for line in posted.json()["sample_receipts"] if line["item_no"] == "SAMPLE-BOX")
        supplier_login(client)
        forbidden = client.post(BASE + f"/internal/receipt-lines/{sample['receipt_line_id']}/link-order",
            json={"factory_id": "huaxing", "order_line_id": paper["id"],
                  "expected_order_revision": 1, "reason": "关联正式订单"})
        assert forbidden.status_code == 403, forbidden.text
        login_as(client, "admin")
        later = _order_payload()
        later.update(contract_no="SAMPLE-1", item_no="SAMPLE-BOX", order_date="2026-09-24",
            due_date="2026-10-01", product_order_quantity="4")
        later["lines"] = [{"packaging_type": "外箱", "paper_quality": "A33+B",
            "specification": "18 × 12.5 × 17.25", "dimension_unit": "cm",
            "usage_quantity": "1", "unit": "个", "unit_price": "2.5"}]
        created = client.post("/api/carton-procurement/orders", json=later)
        assert created.status_code == 201, created.text
        formal = _submit_order(client, created.json())
        target = formal["lines"][0]
        body = {"factory_id": "huaxing", "order_line_id": target["id"],
            "expected_order_revision": formal["revision"], "reason": "样板箱已转正式订单"}
        wrong_unit = client.post(BASE + f"/internal/receipt-lines/{sample['receipt_line_id']}/link-order",
            json={**body, "order_line_id": paper["id"]})
        assert wrong_unit.status_code in (404, 409, 422), wrong_unit.text
        linked = client.post(BASE + f"/internal/receipt-lines/{sample['receipt_line_id']}/link-order", json=body)
        assert linked.status_code == 200, linked.text
        assert linked.json()["order_line_id"] == target["id"]
        duplicate = client.post(BASE + f"/internal/receipt-lines/{sample['receipt_line_id']}/link-order", json=body)
        assert duplicate.status_code == 409, duplicate.text
        with SessionLocal() as db:
            from app.services.carton_replenishment import fulfilled_by_line, posted_receipt_sources
            assert fulfilled_by_line(db, [target["id"]])[target["id"]] == Decimal("4")
            assert posted_receipt_sources(db, [target["id"]])[target["id"]] == {sample["receipt_line_id"]: "4.0000"}
            assert len(db.scalars(select(CartonInventoryMovement).where(
                CartonInventoryMovement.source_id == posted.json()["receipt_id"])).all()) == 2
            receipt = db.get(CartonReceipt, posted.json()["receipt_id"])
            settlement_rows, _, settlement_issues, undated = sources(db, "huaxing", receipt.supplier_id, "2026-09", "CNY")
            assert not settlement_issues and not undated
            assert {row["item_no"] for row in settlement_rows} == {order["item_no"], "SAMPLE-BOX"}
            receipt_revision = receipt.revision
        reversed_receipt = client.post(
            f"/api/carton-procurement/receipts/{posted.json()['receipt_id']}/reverse",
            json={"factory_id": "huaxing", "expected_revision": receipt_revision,
                  "reason": "复核发现整张供应商送货单须冲销"})
        assert reversed_receipt.status_code == 200, reversed_receipt.text
        with SessionLocal() as db:
            assert fulfilled_by_line(db, [target["id"]])[target["id"]] == Decimal("0")
            assert posted_receipt_sources(db, [target["id"]]) == {}
            receipt = db.get(CartonReceipt, posted.json()["receipt_id"])
            settlement_rows, _, _, _ = sources(db, "huaxing", receipt.supplier_id, "2026-09", "CNY")
            assert not settlement_rows


def test_supplier_wrong_delivery_only_keeps_evidence_without_receipt(monkeypatch):
    from openpyxl import Workbook
    with make_client(monkeypatch) as client:
        setup_portal(client)
        workbook = Workbook(); sheet = workbook.active
        sheet.append(["送货单号", "日期", "客户", "客户单号", "客户料号", "名称", "纸质", "规格", "数量"])
        sheet.append(["DK-WRONG-01", "2026-09-24", "华兴", "WRONG-1", "WRONG-BOX", "外箱", "A33+B", "12*10*5 cm", 3])
        buffer = BytesIO(); workbook.save(buffer)
        upload = {"file": ("东康送货明细.xlsx", buffer.getvalue())}
        preview = client.post(BASE + "/shipments/import-preview", files=upload)
        assert preview.status_code == 200 and preview.json()["groups"][0]["ready"] is True
        sent = client.post(BASE + "/shipments/import-confirm", files=upload, data={
            "sha256": preview.json()["sha256"],
            "selections": json.dumps([{"factory_id": "huaxing", "delivery_note_no": "DK-WRONG-01"}])})
        assert sent.status_code == 200, sent.text
        shipment = sent.json()["shipments"][0]
        login_as(client, "warehouse_keeper")
        payload = {"factory_id": "huaxing", "request_id": str(uuid4()), "expected_revision": shipment["revision"],
            "acceptance_date": "2026-09-24", "lines": [{"shipment_line_id": shipment["lines"][0]["id"],
                "received_quantity": 3, "rejected_quantity": 3, "unit_price": 0,
                "no_order_decision": "WRONG_DELIVERY", "difference_reason": "供应商送错纸品"}]}
        confirmed = client.post(BASE + f"/internal/shipments/{shipment['id']}/receive", json=payload)
        assert confirmed.status_code == 200, confirmed.text
        assert confirmed.json()["status"] == "RECEIVED" and confirmed.json()["receipt_id"] is None
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonReceipt, CartonInventoryMovement
        from sqlalchemy import select
        with SessionLocal() as db:
            assert db.scalar(select(CartonReceipt.id).where(CartonReceipt.delivery_note_no == "DK-WRONG-01")) is None
            assert db.scalar(select(CartonInventoryMovement.id).where(CartonInventoryMovement.document_no == "DK-WRONG-01")) is None


def test_supplier_delivery_import_requires_review_then_routes_to_factory_receipt(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client)
        pending = client.get(BASE + "/workspace", params={"factory_id": "huaxing"}).json()["orders"][0]
        content = _dongkang_delivery_file(pending)
        upload = {"file": ("送货明细表.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        preview = client.post(BASE + "/shipments/import-preview", files=upload)
        assert preview.status_code == 200, preview.text
        assert preview.json()["groups"][0]["ready"] is False
        assert "接单" in preview.json()["groups"][0]["rows"][0]["reason"]

        accepted = accept_all(client)
        preview = client.post(BASE + "/shipments/import-preview", files=upload)
        assert preview.status_code == 200, preview.text
        data = preview.json()
        assert data["groups"][0]["ready"] is True
        assert {row["child_no"] for row in data["groups"][0]["rows"]} == {
            line["child_no"] for line in accepted["lines"]}
        form = {"sha256": data["sha256"], "selections": json.dumps([
            {"factory_id": "huaxing", "delivery_note_no": "DK-IMPORT-01"}])}
        confirmed = client.post(BASE + "/shipments/import-confirm", data=form, files=upload)
        assert confirmed.status_code == 200, confirmed.text
        shipment = confirmed.json()["shipments"][0]
        assert shipment["delivery_note_no"] == "DK-IMPORT-01"
        assert shipment["source_filename"] == "送货明细表.xlsx"
        assert shipment["source_sha256"] == data["sha256"]
        assert shipment["status"] == "SENT"
        assert len(shipment["lines"]) == 2
        retry = client.post(BASE + "/shipments/import-confirm", data=form, files=upload)
        assert retry.status_code == 200, retry.text
        assert retry.json()["shipments"][0]["id"] == shipment["id"]
        notification_id = f"carton-shipment:{shipment['id']}"
        assert all(item["id"] != notification_id for item in client.get("/api/system/notifications").json())
        assert client.patch(f"/api/system/notifications/{notification_id}", json={"status": "read"}).status_code == 403
        documents = client.get(BASE + "/documents", params={"factory_id": "huaxing"}).json()
        assert any(row["kind"] == "DELIVERY" and row["document_no"] == "DK-IMPORT-01" for row in documents)

        # Existing in-transit notes from before this notification feature remain actionable.
        from app.db import SessionLocal
        from app.models.auth import SystemNotification
        with SessionLocal() as db:
            db.delete(db.get(SystemNotification, notification_id))
            db.commit()

        login_as(client, "qc_inspector")
        assert all(item["id"] != notification_id for item in client.get("/api/system/notifications").json())
        login_as(client, "a_warehouse_keeper")
        assert all(item["id"] != notification_id for item in client.get("/api/system/notifications").json())
        assert client.patch(f"/api/system/notifications/{notification_id}", json={"status": "read"}).status_code == 404
        login_as(client, "warehouse_keeper")
        notifications = [item for item in client.get("/api/system/notifications").json()
                         if item["id"] == notification_id]
        assert len(notifications) == 1
        assert notifications[0]["status"] == "unread"
        assert notifications[0]["target_factory_id"] == "huaxing"
        assert notifications[0]["payload"]["shipment_id"] == shipment["id"]
        assert client.patch(f"/api/system/notifications/{notification_id}", json={"status": "read"}).status_code == 200
        assert client.patch(f"/api/system/notifications/{notification_id}", json={"status": "handled"}).status_code == 409
        internal = client.get(BASE + "/internal/workspace", params={"factory_id": "huaxing"})
        assert internal.status_code == 200, internal.text
        assert any(row["id"] == shipment["id"] for row in internal.json()["shipments"])
        payload = receive_payload(client, shipment)
        payload["lines"][0]["damaged_quantity"] = 1
        payload["lines"][0]["location_allocations"][0]["quantity"] = 7
        received = client.post(BASE + f"/internal/shipments/{shipment['id']}/receive", json=payload)
        assert received.status_code == 200, received.text
        assert received.json()["status"] == "RECEIVED"
        assert float(received.json()["acceptance_lines"][0]["damaged_quantity"]) == 1
        notification = next(item for item in client.get("/api/system/notifications").json()
                            if item["id"] == notification_id)
        assert notification["status"] == "handled"


def test_supplier_delivery_import_blocks_unmatched_paper_and_other_factory(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client)
        order = accept_all(client)
        for content in (_dongkang_delivery_file(order, missing_item=True),
                        _dongkang_delivery_file(order, destination="华康B"),
                        _dongkang_delivery_file(order, source_spec="31.5 x 11.125 x 11.25 cm")):
            upload = {"file": ("送货明细表.xlsx", content)}
            preview = client.post(BASE + "/shipments/import-preview", files=upload)
            assert preview.status_code == 200, preview.text
            group = preview.json()["groups"][0]
            assert group["ready"] is False
            form = {"sha256": preview.json()["sha256"], "selections": json.dumps([
                {"factory_id": group["factory_id"], "delivery_note_no": group["delivery_note_no"]}])}
            rejected = client.post(BASE + "/shipments/import-confirm", data=form, files=upload)
            assert rejected.status_code == 409, rejected.text
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(_dongkang_delivery_file(order)))
        workbook.active.cell(row=501, column=1, value="TRUNCATED-NOTE")
        output = BytesIO(); workbook.save(output)
        too_long = client.post(BASE + "/shipments/import-preview", files={"file": ("送货明细表.xlsx", output.getvalue())})
        assert too_long.status_code == 422, too_long.text
        assert client.get(BASE + "/workspace", params={"factory_id": "huaxing"}).json()["shipments"] == []


def test_supplier_delivery_import_rolls_back_other_notes_when_one_is_invalid(monkeypatch):
    from openpyxl import load_workbook
    with make_client(monkeypatch) as client:
        setup_portal(client)
        order = accept_all(client)
        workbook = load_workbook(BytesIO(_dongkang_delivery_file(order)))
        sheet = workbook.active
        sheet["A3"] = "DK-IMPORT-02"
        sheet["E3"] = ""
        output = BytesIO(); workbook.save(output)
        upload = {"file": ("送货明细表.xlsx", output.getvalue())}
        preview = client.post(BASE + "/shipments/import-preview", files=upload)
        assert preview.status_code == 200, preview.text
        groups = preview.json()["groups"]
        assert [group["ready"] for group in groups] == [True, False]
        form = {"sha256": preview.json()["sha256"], "selections": json.dumps([
            {"factory_id": "huaxing", "delivery_note_no": group["delivery_note_no"]} for group in groups])}
        rejected = client.post(BASE + "/shipments/import-confirm", data=form, files=upload)
        assert rejected.status_code == 409, rejected.text
        assert client.get(BASE + "/workspace", params={"factory_id": "huaxing"}).json()["shipments"] == []

def test_supplier_mark_templates_require_issued_order_membership_and_ready_latest_version(monkeypatch):
    with make_client(monkeypatch) as client:
        order = setup_portal(client)
        from app.db import SessionLocal
        from app.models.carton_mark import CartonMarkDocument, CartonMarkTemplate
        from app.models.carton_procurement import CartonOrder

        def add_template(db, key, *, factory="huaxing", customer="Dickie", item=None,
                         version=1, status="核对通过", released=False, archived=False):
            template = CartonMarkTemplate(id=key, factory_id=factory,
                customer_name=customer, po=order["contract_no"],
                item=item or order["item_no"], contract_number=order["contract_no"],
                business_key_sha256=key.ljust(64, "0"), document_fingerprint=key.ljust(64, "1"),
                version=version, check_status=status, check_result_json='{"private_check":"secret"}',
                excel_sha256="a" * 64, pdf_sha256="b" * 64,
                created_by="admin", created_at=f"2026-09-21T00:00:{version:02d}",
                updated_at=f"2026-09-21T00:00:{version:02d}",
                manual_released_at="2026-09-21T01:00:00" if released else "",
                is_archived=archived)
            db.add(template)
            db.flush()
            for kind, filename, content_type, content in (
                ("source_excel", "customer.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"excel-data"),
                ("print_pdf", "print.pdf", "application/pdf", b"%PDF-1.4 test"),
            ):
                db.add(CartonMarkDocument(id=f"{key}-{kind}", template_id=key, factory_id=factory,
                    kind=kind, file_name=filename, content_type=content_type,
                    size_bytes=len(content), sha256="a" * 64, content=content,
                    created_at="2026-09-21T00:00:00"))

        with SessionLocal() as db:
            add_template(db, "match-v1")
            add_template(db, "wrong-customer", customer="Other")
            add_template(db, "wrong-item", item="OTHER-ITEM")
            add_template(db, "wrong-factory", factory="huadeng")
            db.commit()

        url = BASE + "/carton-mark/templates"
        download = BASE + "/carton-mark/templates/{}/documents/{}"
        records = client.get(url, params={"factory_id": "huaxing"})
        assert records.status_code == 200, records.text
        assert [row["id"] for row in records.json()] == ["match-v1"]
        assert "private_check" not in records.text and "created_by" not in records.text
        assert client.get(download.format("match-v1", "source_excel"), params={"factory_id": "huaxing"}).content == b"excel-data"
        assert client.get(download.format("match-v1", "print_pdf"), params={"factory_id": "huaxing"}).content == b"%PDF-1.4 test"
        preview = client.get(download.format("match-v1", "print_pdf"), params={"factory_id": "huaxing", "preview": "true"})
        assert preview.status_code == 200 and preview.headers["content-disposition"].startswith("inline;")
        for hidden in ("wrong-customer", "wrong-item", "wrong-factory"):
            assert client.get(download.format(hidden, "print_pdf"), params={"factory_id": "huaxing"}).status_code == 404
        assert client.get(url, params={"factory_id": "huadeng"}).status_code in (403, 422)
        assert client.get(download.format("match-v1", "other"), params={"factory_id": "huaxing"}).status_code == 422
        assert client.get("/api/carton-mark/templates", params={"factory_id": "huaxing"}).status_code == 403

        with SessionLocal() as db:
            add_template(db, "match-v2", version=2, status="需复核")
            db.commit()
        assert client.get(url, params={"factory_id": "huaxing"}).json() == []
        assert client.get(download.format("match-v1", "print_pdf"), params={"factory_id": "huaxing"}).status_code == 404

        with SessionLocal() as db:
            db.get(CartonMarkTemplate, "match-v2").manual_released_at = "2026-09-21T01:00:00"
            db.commit()
        assert [row["id"] for row in client.get(url, params={"factory_id": "huaxing"}).json()] == ["match-v2"]

        with SessionLocal() as db:
            db.get(CartonMarkTemplate, "match-v2").is_archived = True
            db.commit()
        assert client.get(url, params={"factory_id": "huaxing"}).json() == []

        with SessionLocal() as db:
            db.get(CartonMarkTemplate, "match-v2").is_archived = False
            db.get(CartonOrder, order["id"]).status = "CANCELLED"
            db.commit()
        assert client.get(download.format("match-v2", "print_pdf"), params={"factory_id": "huaxing"}).status_code == 404

        with SessionLocal() as db:
            db.get(CartonOrder, order["id"]).status = "PENDING_SUPPLIER"
            db.commit()
        revoke_supplier_permission()
        assert client.get(url, params={"factory_id": "huaxing"}).status_code == 403


def test_supplier_permission_scope_and_whitelist(monkeypatch):
    with make_client(monkeypatch) as client:
        assert client.get(BASE+"/workspace", params={"factory_id":"huaxing"}).status_code == 401
        order = setup_portal(client)
        workspace = client.get(BASE+"/workspace", params={"factory_id":"huaxing"})
        assert workspace.status_code == 200
        serialized = json.dumps(workspace.json())
        for forbidden in ("unit_price", "currency", "note", "created_by", "customer_code", "location_allocations", "supplier_id"):
            assert f'"{forbidden}"' not in serialized
        assert client.get(BASE+"/workspace", params={"factory_id":"huadeng"}).status_code in (403,422)
        assert client.get(BASE+"/internal/workspace", params={"factory_id":"huaxing"}).status_code == 403
        assert client.get("/api/carton-procurement/orders",params={"factory_id":"huaxing"}).status_code == 403
        assert client.put(BASE+"/internal/members", json={"factory_id":"huaxing","username":"supplier-test","expected_revision":1,"reason":"旧绑定接口应停用"}).status_code == 404
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonOrder, CartonSupplier
        with SessionLocal() as db:
            db.add(CartonSupplier(id="other-supplier",factory_id="huaxing",supplier_code="OTHER",supplier_name="其他供应商",status="ACTIVE",created_at="2026-09-21",updated_at="2026-09-21")); db.flush()
            db.get(CartonOrder,order["id"]).supplier_id="other-supplier"; db.commit()
        assert client.get(BASE+"/workspace",params={"factory_id":"huaxing"}).status_code == 403
        assert client.put(BASE+f"/papers/{order['lines'][0]['id']}/commitment",json={"factory_id":"huaxing","issue_id":"guess","expected_revision":0,"promised_date":"2026-09-23"}).status_code == 403
        login_as(client,"warehouse_keeper")
        assert client.put(BASE+"/internal/members",json={"factory_id":"huaxing","username":"supplier-test","expected_revision":1,"reason":"旧绑定接口应停用"}).status_code == 404
        assert client.get(BASE+"/workspace",params={"factory_id":"huaxing"}).status_code == 403


def test_supplier_read_edit_approve_are_independent(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client)
        revoke_supplier_permission("carton_supplier:edit")
        revoke_supplier_permission("carton_supplier:approve")
        supplier_login(client)
        workspace = client.get(BASE + "/workspace", params={"factory_id": "huaxing"})
        assert workspace.status_code == 200
        assert client.get(BASE + "/documents", params={"factory_id": "huaxing"}).status_code == 200
        order = workspace.json()["orders"][0]
        line = order["lines"][0]
        commitment = {"factory_id": "huaxing", "issue_id": order["issue_id"],
            "expected_revision": line["commitment_revision"], "promised_date": "2026-09-23"}
        assert client.put(BASE + f"/papers/{line['id']}/commitment", json=commitment).status_code == 403
        assert client.post(BASE + "/shipments", json=ship_payload(order)).status_code == 403

        restore_supplier_permission("carton_supplier:approve")
        supplier_login(client)
        accepted_order = accept_all(client)
        assert client.post(BASE + "/shipments", json=ship_payload(accepted_order)).status_code == 403

        restore_supplier_permission("carton_supplier:edit")
        supplier_login(client)
        assert client.post(BASE + "/shipments", json=ship_payload(accepted_order)).status_code == 201
        revoke_supplier_permission("carton_supplier:read")
        supplier_login(client, expect_read=False)
        assert client.get(BASE + "/workspace", params={"factory_id": "huaxing"}).status_code == 403
        assert client.put(BASE + f"/papers/{line['id']}/commitment", json=commitment).status_code == 403


def test_supplier_permission_can_be_revoked_from_iam_without_changing_internal_roles(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client)
        login_as(client, "admin")
        access = client.get("/api/iam/users/supplier-test/access")
        assert access.status_code == 200, access.text
        preview = client.post("/api/iam/users/supplier-test/access/preview", json={
            "base_revision": access.json()["authorization_version"],
            "reason": "验证供应商送货权限单独撤回",
            "overrides": [{"permission_code": "carton_supplier:edit", "effect": "inherit",
                "factory_id": "*", "department": "*"}],
        })
        assert preview.status_code == 200, preview.text
        committed = client.post("/api/iam/users/supplier-test/access/commit", json={
            "preview_token": preview.json()["preview_token"], "confirm_high_risk": True,
        })
        assert committed.status_code == 200, committed.text
        supplier_login(client)
        assert client.get(BASE + "/workspace", params={"factory_id": "huaxing"}).status_code == 200
        assert client.post(BASE + "/shipments", json={"factory_id": "huaxing",
            "request_id": str(uuid4()), "delivery_note_no": "DENIED", "delivery_date": "2026-09-24",
            "lines": [{"order_line_id": "missing", "issue_id": "missing", "quantity": 1}]}).status_code == 403


def test_supplier_permissions_do_not_require_global_auth_mode_migration(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client)
        from app.db import SessionLocal
        from app.services.auth import ensure_authz_startup_safety
        with SessionLocal() as db:
            ensure_authz_startup_safety(db)


def test_supplier_account_sees_dongkang_issued_factories_without_internal_factory_role(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client)
        supplier_login(client)
        memberships = client.get(BASE + "/memberships")
        assert memberships.status_code == 200
        assert [item["factory_id"] for item in memberships.json()] == ["huaxing"]
        assert client.get(BASE + "/workspace", params={"factory_id": "huakang-a"}).status_code == 403

        login_as(client, "admin")
        customer = client.post("/api/carton-procurement/customers", json={
            "factory_id": "huakang-a", "customer_code": "DICKIE", "customer_name": "Dickie", "country_region": "德国"})
        assert customer.status_code == 201, customer.text
        payload = _order_payload()
        payload.update(factory_id="huakang-a", contract_no="SC-HUAKANG-A")
        created = client.post("/api/carton-procurement/orders", json=payload)
        assert created.status_code == 201, created.text
        order = created.json()
        submitted = client.post(f"/api/carton-procurement/orders/{order['order_no']}/submit-supplier",
            json={"factory_id": "huakang-a", "expected_revision": order["revision"]})
        assert submitted.status_code == 200, submitted.text
        issued = client.post(f"/api/carton-procurement/orders/{order['order_no']}/purchase-order-issues.xlsx",
            json={"factory_id": "huakang-a", "expected_revision": submitted.json()["revision"]})
        assert issued.status_code == 200, issued.text

        supplier_login(client)
        memberships = client.get(BASE + "/memberships")
        assert [item["factory_id"] for item in memberships.json()] == ["huakang-a", "huaxing"]
        scope = client.get(BASE + "/workspace", params={"factory_id": "huakang-a"})
        assert scope.status_code == 200, scope.text
        assert [item["contract_no"] for item in scope.json()["orders"]] == ["SC-HUAKANG-A"]
        assert client.get(BASE + "/carton-mark/templates", params={"factory_id": "huakang-a"}).status_code == 200
        supplier_order = scope.json()["orders"][0]
        assert supplier_order["order_date"] == order["order_date"]
        paper = supplier_order["lines"][0]
        accepted = client.put(BASE + "/commitments/batch", json={
            "factory_id": "huakang-a", "lines": [{
                "order_line_id": line["id"], "issue_id": supplier_order["issue_id"],
                "expected_revision": line["commitment_revision"], "promised_date": "2026-09-23",
            } for line in supplier_order["lines"]]})
        assert accepted.status_code == 200, accepted.text
        shipment = client.post(BASE + "/shipments", json={
            "factory_id": "huakang-a", "request_id": str(uuid4()),
            "delivery_note_no": "DN-HUAKANG-A", "delivery_date": "2026-09-23",
            "lines": [{"order_line_id": paper["id"], "issue_id": supplier_order["issue_id"], "quantity": 1}]})
        assert shipment.status_code == 201, shipment.text
        assert client.get("/api/carton-procurement/orders", params={"factory_id": "huakang-a"}).status_code == 403
        assert client.get(BASE + "/workspace", params={"factory_id": "huadeng"}).status_code == 403

        revoke_supplier_permission()
        supplier_login(client, expect_read=False)
        assert client.get(BASE + "/memberships").json() == []
        assert client.get(BASE + "/workspace", params={"factory_id": "huakang-a"}).status_code == 403


def test_confirmed_order_is_immediately_visible_to_supplier_without_separate_issue(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client)
        login_as(client, "admin")
        created = _create_order(client, submit_supplier=False)
        confirmed = client.post(f"/api/carton-procurement/orders/{created['order_no']}/submit-supplier",
            json={"factory_id": "huaxing", "expected_revision": created["revision"]})
        assert confirmed.status_code == 200, confirmed.text
        context = client.get(f"/api/carton-procurement/orders/{created['order_no']}/purchase-order-context",
            params={"factory_id": "huaxing"}).json()
        assert context["pending_type"] == "NONE"
        assert len(context["issues"]) == 1
        assert context["issues"][0]["document_no"].endswith("-P00")
        redownload = client.post(f"/api/carton-procurement/orders/{created['order_no']}/purchase-order-issues.xlsx",
            json={"factory_id": "huaxing", "expected_revision": confirmed.json()["revision"]})
        assert redownload.status_code == 200, redownload.text
        assert redownload.headers["x-purchase-order-issue-id"] == context["issues"][0]["id"]
        supplier_login(client)
        orders = client.get(BASE + "/workspace", params={"factory_id": "huaxing"}).json()["orders"]
        assert created["order_no"] in {row["order_no"] for row in orders}

def test_partial_receipt_atomicity_replay_and_remaining(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client); order=accept_all(client); payload=ship_payload(order)
        shipped=client.post(BASE+"/shipments",json=payload)
        assert shipped.status_code == 201, shipped.text
        shipment=shipped.json()
        assert client.post(BASE+"/shipments",json=payload).json()["id"] == shipment["id"]
        altered={**payload,"delivery_note_no":"changed"}
        assert client.post(BASE+"/shipments",json=altered).status_code == 409
        login_as(client,"warehouse_keeper")
        assert client.get("/api/carton-procurement/receipts",params={"factory_id":"huaxing"}).json()["total"] == 0
        receive=receive_payload(client,shipment)
        excess=json.loads(json.dumps(receive));excess["lines"][0]["received_quantity"]=11
        assert client.post(BASE+f"/internal/shipments/{shipment['id']}/receive",json=excess).status_code == 422
        invalid=json.loads(json.dumps(receive));invalid["lines"][-1]["unit_price"]=0
        denied=client.post(BASE+f"/internal/shipments/{shipment['id']}/receive",json=invalid)
        assert denied.status_code == 422,denied.text
        assert client.get("/api/carton-procurement/receipts",params={"factory_id":"huaxing"}).json()["total"] == 0
        received=client.post(BASE+f"/internal/shipments/{shipment['id']}/receive",json=receive)
        assert received.status_code == 200,received.text
        replay=client.post(BASE+f"/internal/shipments/{shipment['id']}/receive",json=receive)
        assert replay.status_code == 200 and replay.json()["receipt_id"] == received.json()["receipt_id"]
        receive["request_id"]=str(uuid4())
        assert client.post(BASE+f"/internal/shipments/{shipment['id']}/receive",json=receive).status_code == 409
        assert client.get("/api/carton-procurement/receipts",params={"factory_id":"huaxing"}).json()["total"] == 1
        supplier_login(client)
        data=client.get(BASE+"/workspace",params={"factory_id":"huaxing"}).json()
        assert data["orders"][0]["status"] == "PARTIALLY_RECEIVED"
        assert [float(line["received_quantity"]) for line in data["orders"][0]["lines"]] == [8,8]
        assert float(data["orders"][0]["lines"][0]["remaining_to_ship"]) == 22
        assert "unit_price" not in json.dumps(data["shipments"])
        next_delivery=ship_payload(order);next_delivery["delivery_note_no"]="DN-NEXT-PARTIAL"
        response=client.post(BASE+"/shipments",json=next_delivery)
        assert response.status_code == 201,response.text
        login_as(client,"warehouse_keeper")
        current=client.get("/api/carton-procurement/orders",params={"factory_id":"huaxing"}).json()["items"][0]
        returned=client.post(f"/api/carton-procurement/orders/{current['order_no']}/return",json={"factory_id":"huaxing","expected_revision":current["revision"],"reason":"已有在途不能先退单"})
        assert returned.status_code == 409 and "在途" in returned.text,returned.text

def test_positive_receipts_require_complete_locations_before_any_write(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client); order = accept_all(client)
        shipment = client.post(BASE+"/shipments", json=ship_payload(order)).json()
        login_as(client, "warehouse_keeper")
        payload = receive_payload(client, shipment)
        from sqlalchemy import func, select
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonReceipt, CartonInventoryMovement
        from app.models.carton_supplier_portal import SupplierShipment
        for allocations, message in [([], "必须选择"), ([{**payload["lines"][-1]["location_allocations"][0], "quantity": 7}], "之和必须等于"), ([{"location_id": "missing-location", "quantity": 8}], "仓位")]:
            invalid = json.loads(json.dumps(payload))
            invalid["lines"][-1]["location_allocations"] = allocations
            result = client.post(BASE+f"/internal/shipments/{shipment['id']}/receive", json=invalid)
            assert result.status_code in (404, 422) and message in result.text, result.text
            with SessionLocal() as db:
                assert db.scalar(select(func.count()).select_from(CartonReceipt)) == 0
                assert db.scalar(select(func.count()).select_from(CartonInventoryMovement)) == 0
                assert db.get(SupplierShipment, shipment["id"]).status == "SENT"
        # Physically received but fully rejected paper has zero effective stock and needs no allocation.
        payload["lines"][-1].update(received_quantity=10, rejected_quantity=10, location_allocations=[], difference_reason="全部拒收不计入库存")
        result = client.post(BASE+f"/internal/shipments/{shipment['id']}/receive", json=payload)
        assert result.status_code == 200, result.text
        with SessionLocal() as db:
            movements = list(db.scalars(select(CartonInventoryMovement)).all())
            assert len(movements) == 1 and float(movements[0].quantity) == 8
            assert db.get(SupplierShipment, shipment["id"]).status == "RECEIVED"


@pytest.mark.parametrize("all_zero", [False, True])
def test_unreceived_papers_release_transit_without_positive_inventory(monkeypatch, all_zero):
    with make_client(monkeypatch) as client:
        raw = setup_portal(client); order = accept_all(client)
        shipment = client.post(BASE+"/shipments", json=ship_payload(order)).json()
        login_as(client,"warehouse_keeper")
        payload = receive_payload(client, shipment)
        zeros = payload["lines"] if all_zero else payload["lines"][-1:]
        for line in zeros:
            line["received_quantity"] = 0; line["unit_price"] = 0; line["location_allocations"] = []
            line["difference_reason"] = "本纸品尚未实际送到"
        if all_zero:
            from app.db import SessionLocal
            from app.models.carton_procurement import CartonOrder
            with SessionLocal() as db:
                db.get(CartonOrder,raw["id"]).status="CANCELLED";db.commit()
        result=client.post(BASE+f"/internal/shipments/{shipment['id']}/receive",json=payload)
        assert result.status_code == 200,result.text
        assert result.json()["status"] == ("NOT_RECEIVED" if all_zero else "RECEIVED")
        notice = next(item for item in client.get("/api/system/notifications").json()
                      if item["id"] == f"carton-shipment:{shipment['id']}")
        assert notice["status"] == "handled"
        receipts=client.get("/api/carton-procurement/receipts",params={"factory_id":"huaxing"}).json()
        assert receipts["total"] == (0 if all_zero else 1)
        if not all_zero:
            assert len(receipts["items"][0]["lines"]) == 1
        repeated=client.post(BASE+f"/internal/shipments/{shipment['id']}/receive",json=payload)
        assert repeated.status_code == 200 and repeated.json()["id"] == shipment["id"]
        supplier_login(client)
        result=client.get(BASE+"/workspace",params={"factory_id":"huaxing"}).json()
        assert all(float(line["in_transit_quantity"]) == 0 for line in result["orders"][0]["lines"])
        assert float(result["orders"][0]["lines"][-1]["received_quantity"]) == 0


def test_concurrent_shipments_reserve_capacity_and_replay(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client);order=accept_all(client);payload=ship_payload(order)
        payload["lines"]=payload["lines"][:1];payload["lines"][0]["quantity"]=20
        from app.db import SessionLocal
        from app.models.auth import AuthUser
        from app.services.auth import build_auth_context
        from app.services.carton_supplier_portal import create_shipment
        from app.schemas.carton_supplier_portal import ShipmentCreate
        barrier=Barrier(2)
        def execute(body):
            with SessionLocal() as db:
                user=build_auth_context(db,db.get(AuthUser,"supplier-test"));barrier.wait()
                try:return create_shipment(db,user,ShipmentCreate(**body))["id"]
                except Exception as exc:db.rollback();return getattr(exc,"status_code",str(exc))
        with ThreadPoolExecutor(max_workers=2) as pool:
            futures=[pool.submit(execute,payload) for _ in range(2)]
            results=[future.result(timeout=30) for future in futures]
        assert results[0] == results[1] and isinstance(results[0],str),results
        second={**payload,"request_id":str(uuid4()),"delivery_note_no":"SECOND"}
        assert client.post(BASE+"/shipments",json=second).status_code == 409
        second["lines"][0]["quantity"]=10
        assert client.post(BASE+"/shipments",json=second).status_code == 201

def test_cross_entry_transit_blocks_reduction_cancellation_and_manual_receipt(monkeypatch):
    with make_client(monkeypatch) as client:
        raw=setup_portal(client);order=accept_all(client)
        shipment=client.post(BASE+"/shipments",json=ship_payload(order)).json()
        login_as(client,"warehouse_keeper")
        for action,payload in [("reduce", {"reduction_quantity":3600}), ("cancel", {})]:
            result=client.post(f"/api/carton-procurement/orders/{raw['order_no']}/{action}",json={"factory_id":"huaxing","expected_revision":raw["revision"],"reason":"测试跨入口需求保护",**payload})
            assert result.status_code == 409 and "在途" in result.text,result.text
        bulk=client.post("/api/carton-procurement/orders/bulk-cancel",json={"factory_id":"huaxing","reason":"测试批量取消跨入口", "items":[{"order_no":raw["order_no"],"expected_revision":raw["revision"]}]})
        assert bulk.status_code == 409 and "在途" in bulk.text,bulk.text
        for note in (shipment["delivery_note_no"],"DIFFERENT-NOTE"):
            response=client.post("/api/carton-procurement/receipts",json={"factory_id":"huaxing","delivery_note_no":note,"delivery_date":"2026-09-21", "lines":[{"order_line_id":raw["lines"][0]["id"],"delivered_quantity":1,"received_quantity":1,"unit_price":2}]})
            assert response.status_code == 409 and "供应商" in response.text,response.text
        receive=receive_payload(client,shipment)
        for line in receive["lines"]:
            line.update(received_quantity=0,unit_price=0,location_allocations=[],difference_reason="供应商误报尚未送达")
        assert client.post(BASE+f"/internal/shipments/{shipment['id']}/receive",json=receive).status_code == 200
        reduced=client.post(f"/api/carton-procurement/orders/{raw['order_no']}/reduce",json={"factory_id":"huaxing","expected_revision":raw["revision"],"reason":"未发货允许全量退单","reduction_quantity":3600})
        assert reduced.status_code == 200 and reduced.json()["status"] == "CANCELLED",reduced.text


def test_supplier_cannot_reuse_an_existing_receipt_delivery_note(monkeypatch):
    with make_client(monkeypatch) as client:
        raw=setup_portal(client);order=accept_all(client)
        login_as(client,"warehouse_keeper")
        existing=client.post("/api/carton-procurement/receipts",json={"factory_id":"huaxing","delivery_note_no":"DN-PORTAL","delivery_date":"2026-09-21","lines":[{"order_line_id":raw["lines"][0]["id"],"delivered_quantity":1,"received_quantity":1,"unit_price":2}]})
        assert existing.status_code == 201,existing.text
        supplier_login(client)
        response=client.post(BASE+"/shipments",json=ship_payload(order))
        assert response.status_code == 409 and "已有收料记录" in response.text,response.text


def test_state_changes_unissued_versions_and_permission_revocation(monkeypatch):
    with make_client(monkeypatch) as client:
        raw=setup_portal(client);order=accept_all(client)
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonOrder,CartonOrderLine
        with SessionLocal() as db:
            db.get(CartonOrderLine,order["lines"][0]["id"]).required_quantity+=1;db.commit()
        data=client.get(BASE+"/workspace",params={"factory_id":"huaxing"}).json()["orders"][0]
        assert data["awaiting_issue"] and not data["lines"][0]["accepted"]
        assert float(data["lines"][0]["required_quantity"]) == 30
        assert client.post(BASE+"/shipments",json=ship_payload(order)).status_code == 409
        with SessionLocal() as db:
            db.get(CartonOrder,raw["id"]).status="CANCELLED";db.commit()
        assert client.put(BASE+f"/papers/{order['lines'][0]['id']}/commitment",json={"factory_id":"huaxing","issue_id":order["issue_id"],"expected_revision":1,"promised_date":"2026-09-24"}).status_code == 409
        revoke_supplier_permission()
        supplier_login(client, expect_read=False)
        assert client.get(BASE+"/workspace",params={"factory_id":"huaxing"}).status_code == 403

def test_replacement_sources_cannot_be_misclassified_as_ordinary_supplier_delivery(monkeypatch):
    with make_client(monkeypatch) as client:
        raw=setup_portal(client);order=accept_all(client)
        import app.services.carton_supplier_portal as portal
        monkeypatch.setattr(portal,"blocked_replacement_lines",lambda db,factory,ids:{order["lines"][0]["id"]})
        result=client.get(BASE+"/workspace",params={"factory_id":"huaxing"}).json()["orders"][0]
        assert "补单" in result["lines"][0]["shipping_blocked_reason"]
        response=client.post(BASE+"/shipments",json=ship_payload(order))
        assert response.status_code == 409 and "补单" in response.text


def test_internal_staff_can_hold_supplier_permissions_without_changing_internal_scope(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client)
        from app.db import SessionLocal
        from app.models.auth import AuthRole, AuthUser, AuthUserRole
        from app.services.auth import build_auth_context, has_permission_in_scope
        from sqlalchemy import select

        with SessionLocal() as db:
            role = db.scalar(select(AuthRole).where(AuthRole.code == "position_qc_inspector"))
            assert role is not None
            db.add(AuthUserRole(id="supplier-now-internal", user_id="supplier-test", role_id=role.id,
                                factory_id="huakang-b", department="qc"))
            db.commit()
            context = build_auth_context(db, db.get(AuthUser, "supplier-test"))
            assert has_permission_in_scope(context, "qc_inspection:read", "huakang-b", "qc")
            assert not has_permission_in_scope(context, "qc_inspection:read", "huaxing", "qc")

        login_as(client, "admin")
        access = client.get("/api/iam/users/supplier-test/access")
        assert access.status_code == 200, access.text
        assert access.json()["cleanup_override_count"] == 0

        login = client.post("/api/auth/login", json={"username": "supplier-test", "password": PASSWORD})
        assert login.status_code == 200, login.text
        assert "qc_inspection:read" in login.json()["permissions"]
        assert [item["factory_id"] for item in client.get(BASE + "/memberships").json()] == ["huaxing"]
        assert client.get(BASE + "/workspace", params={"factory_id": "huaxing"}).status_code == 200
        assert client.get(BASE + "/carton-mark/templates", params={"factory_id": "huaxing"}).status_code == 200
        assert client.get("/api/carton-procurement/orders", params={"factory_id": "huaxing"}).status_code == 403
        assert client.get(BASE + "/workspace", params={"factory_id": "huakang-b"}).status_code == 403

        revoke_supplier_permission()
        login = client.post("/api/auth/login", json={"username": "supplier-test", "password": PASSWORD})
        assert login.status_code == 200, login.text
        assert client.get(BASE + "/memberships").json() == []
        assert client.get(BASE + "/workspace", params={"factory_id": "huaxing"}).status_code == 403
        with SessionLocal() as db:
            context = build_auth_context(db, db.get(AuthUser, "supplier-test"))
            assert has_permission_in_scope(context, "qc_inspection:read", "huakang-b", "qc")


def test_existing_internal_account_can_receive_supplier_permission(monkeypatch):
    with make_client(monkeypatch) as client:
        setup_portal(client)
        from app.db import SessionLocal
        from app.models.auth import AuthUser
        with SessionLocal() as db:
            account = db.query(AuthUser).filter_by(username="warehouse_keeper").one()
            account_id = account.id
        grant_supplier_permissions(user_id=account_id, permissions=("carton_supplier:read",))
        login_as(client, "warehouse_keeper")
        assert [item["factory_id"] for item in client.get(BASE + "/memberships").json()] == ["huaxing"]
        assert client.get(BASE + "/workspace", params={"factory_id": "huaxing"}).status_code == 200
        assert client.get("/api/carton-procurement/orders", params={"factory_id": "huaxing"}).status_code == 200


def test_attachment_formats_version_scope_and_download(monkeypatch):
    with make_client(monkeypatch) as client:
        raw=setup_portal(client)
        from pypdf import PdfWriter
        writer=PdfWriter();writer.add_blank_page(width=100,height=100);buffer=BytesIO();writer.write(buffer);content=buffer.getvalue()
        url=BASE+f"/internal/orders/{raw['id']}/attachments"
        assert client.post(url,data={"factory_id":"huaxing"},files={"file":("drawing.pdf",content)}).status_code == 403
        login_as(client,"admin")
        for filename,data in (("evil.html",b"<html>"),("fake.pdf",b"not pdf"),("../x.pdf",content),("fake.docx",b"not office"),("huge.pdf",b"x"*(10*1024*1024+1))):
            response=client.post(url,data={"factory_id":"huaxing"},files={"file":(filename,data)})
            assert response.status_code == 422,(filename,response.text)
        result=client.post(url,data={"factory_id":"huaxing"},files={"file":("drawing.pdf",content)})
        assert result.status_code == 201,result.text
        attachment=result.json();assert attachment["version"] == 1
        assert client.post(url,data={"factory_id":"huaxing"},files={"file":("drawing.pdf",content)}).json()["id"] == attachment["id"]
        writer.add_blank_page(width=100,height=100);next_buffer=BytesIO();writer.write(next_buffer)
        next_version=client.post(url,data={"factory_id":"huaxing"},files={"file":("drawing.pdf",next_buffer.getvalue())})
        assert next_version.status_code == 201 and next_version.json()["version"] == 2
        assert next_version.json()["sha256"] != attachment["sha256"]
        supplier_login(client)
        download=client.get(BASE+f"/attachments/{attachment['id']}",params={"factory_id":"huaxing"})
        assert download.status_code == 200 and download.content == content
        assert download.headers["content-disposition"].startswith("attachment;")
        assert download.headers["x-content-type-options"] == "nosniff"
        assert client.get(BASE+f"/attachments/{attachment['id']}",params={"factory_id":"huadeng"}).status_code in (403,422)
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonOrder,CartonSupplier
        with SessionLocal() as db:
            db.add(CartonSupplier(id="attachment-other",factory_id="huaxing",supplier_code="OTHER",supplier_name="其他供应商",status="ACTIVE",created_at="2026-09-21",updated_at="2026-09-21"));db.flush()
            db.get(CartonOrder,raw["id"]).supplier_id="attachment-other";db.commit()
        assert client.get(BASE+f"/attachments/{attachment['id']}",params={"factory_id":"huaxing"}).status_code == 403
        with SessionLocal() as db:
            db.get(CartonOrder,raw["id"]).supplier_id=raw["supplier_id"]
            db.get(CartonOrder,raw["id"]).status="CONFIRMED";db.commit()
        assert client.get(BASE+f"/attachments/{attachment['id']}",params={"factory_id":"huaxing"}).status_code == 403

def test_portal_migration_in_isolated_database(monkeypatch,tmp_path):
    with make_client(monkeypatch):
        from app.db import Base
        from app.services.carton_supplier_portal import TABLES
        from sqlalchemy import create_engine,inspect,text
        from alembic.migration import MigrationContext
        from alembic.operations import Operations
        engine=create_engine(f"sqlite:///{tmp_path/'migration.db'}")
        Base.metadata.create_all(engine,tables=[table for name,table in Base.metadata.tables.items() if name not in TABLES])
        path=Path(__file__).parents[1]/"alembic/versions/20260921_0118_carton_supplier_portal.py"
        spec=importlib.util.spec_from_file_location("portal_migration",path);module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
        assert module.down_revision == "20260917_0117"
        with engine.begin() as conn:
            conn.execute(text("CREATE TABLE evidence_test (id INTEGER PRIMARY KEY, value TEXT)"));conn.execute(text("INSERT INTO evidence_test VALUES (1,'preserve')"))
            with Operations.context(MigrationContext.configure(conn)):module.upgrade()
            assert (set(TABLES) - {"carton_supplier_unmatched_lines"}).issubset(inspect(conn).get_table_names())
            next_path=Path(__file__).parents[1]/"alembic/versions/20260924_0123_supplier_unmatched_delivery.py"
            next_spec=importlib.util.spec_from_file_location("unmatched_migration",next_path)
            next_module=importlib.util.module_from_spec(next_spec);next_spec.loader.exec_module(next_module)
            assert next_module.down_revision == "20260924_0122"
            with Operations.context(MigrationContext.configure(conn)):next_module.upgrade()
            assert set(TABLES).issubset(inspect(conn).get_table_names())
            assert conn.execute(text("SELECT value FROM evidence_test")).scalar() == "preserve"
            assert conn.exec_driver_sql("PRAGMA foreign_key_check").fetchall() == []
        engine.dispose()
