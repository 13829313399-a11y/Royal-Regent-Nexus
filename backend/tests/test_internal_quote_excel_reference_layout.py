import json
from io import BytesIO
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
from PIL import Image

from app.services.internal_quote_excel import _build_summary_sheet


def _find_cell(sheet, value: object):
    return next(
        cell
        for row in sheet.iter_rows()
        for cell in row
        if cell.value == value
    )


def _summary_sheet(
    *,
    with_image: bool = False,
    customer: str = "测试客户",
    sales_payload: dict | None = None,
    route_rows: list[dict] | None = None,
    markup_tiers: list[dict] | None = None,
    freight_enabled: bool = True,
    lifting_enabled: bool = True,
):
    workbook = Workbook()
    quote = SimpleNamespace(
        product_name="参考公式与主图测试",
        quote_no="IQ-REFERENCE-LAYOUT",
        remark="产品功能介绍",
        customer=customer,
    )
    summary = {
        "t1": [
            {"key": "material", "value": "1.25"},
            {"key": "sewing_hair", "value": "2.5"},
            {"key": "sewing_cloth", "value": "6.5"},
            {"key": "hardware", "value": "0.75"},
            {"key": "electronic", "value": "0.8"},
            {"key": "motor", "value": "0.3"},
            {"key": "suction", "value": "0.4"},
            {"key": "glue_bag", "value": "0.2"},
        ],
        "t2": [
            {"key": "color_box", "value": "0.45"},
            {"key": "battery", "value": "0.15"},
            {"key": "libao", "value": "0.16"},
            {"key": "plating", "value": "0.17"},
            {"key": "other_buy", "value": "0.18"},
            {"key": "carton", "value": "0.55"},
        ],
        "t3": [
            {"key": "injection_labor", "value": "0.7"},
            {"key": "painting_labor", "value": "0.2"},
            {"key": "paint_material", "value": "0.1"},
        ],
        "t4": [
            {
                "key": "sewcloth13",
                "label": "车衣物料退税（仅华康C/D）",
                "amount_hkd": "2.0",
                "rate_percent": "11.5",
            }
        ],
        "shipping_pricing": {
            "markup": "1.2",
            "active_markup_moq": "10000",
            "markup_tiers": markup_tiers or [
                {"moq": 3000, "markup": "1.2"},
                {"moq": 5000, "markup": "1.2"},
                {"moq": 10000, "markup": "1.1", "is_active": True},
            ],
            "misc_ratio": "0.02",
            "freight_enabled": freight_enabled,
            "lifting_enabled": lifting_enabled,
            "rows": route_rows or [],
        },
        "molding_material_breakdown": {
            "total_hkd": "1.25",
            "imported_hkd": "0.75",
            "domestic_hkd": "0.5",
        },
    }
    attachments = []
    if with_image:
        buffer = BytesIO()
        Image.new("RGB", (1200, 400), (16, 118, 110)).save(buffer, "PNG")
        attachments.append(
            SimpleNamespace(
                id="main-image",
                department="product-image",
                content_type="image/png",
                content=buffer.getvalue(),
                uploaded_at="2026-08-20T10:00:00",
            )
        )
    sections = [] if sales_payload is None else [
        SimpleNamespace(
            department="sales",
            payload_json=json.dumps(sales_payload, ensure_ascii=False),
            calculation_json=json.dumps({}, ensure_ascii=False),
        )
    ]
    _build_summary_sheet(
        workbook,
        quote,
        sections,
        {"fx": {"rmb_hkd": "0.85", "hkd_usd": "7.8"}},
        summary,
        {
            "molding_hkd": 1.95,
            "painting_hkd": 0.3,
            "sewing_hkd": 9.0,
        },
        attachments,
    )
    return workbook, workbook.active


def test_selective_output_compacts_routes_moqs_and_removes_disabled_testing_box():
    workbook, sheet = _summary_sheet(
        sales_payload={"testing_fee_enabled": False},
        route_rows=[
            {"name": "出厂价", "freight_hkd": "0", "lift_hkd": "0"},
            {"name": "HK 40 柜", "freight_hkd": "1.20", "lift_hkd": "0.50"},
            {"name": "YT 5 吨车", "freight_hkd": "2.30", "lift_hkd": "0.80"},
        ],
        markup_tiers=[
            {"moq": 3000, "markup": "1.25", "include_in_output": False},
            {"moq": 5000, "markup": "1.20", "include_in_output": True},
            {"moq": 10000, "markup": "1.10", "is_active": True, "include_in_output": True},
        ],
    )
    try:
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        assert "测试费用" not in values
        assert "包含测试费用（US）：" not in values
        assert "报价（MOQ3K）" not in values
        assert "报价（MOQ5K）" in values
        assert "报价（MOQ10K）" in values

        route_header = _find_cell(sheet, "运输方案").row
        assert sheet.cell(route_header, 5).value == "HK 40 柜"
        assert sheet.cell(route_header, 6).value == "YT 5 吨车"
        assert sheet.cell(route_header, 7).value is None
        assert "HK 20 柜" not in values
        subtotal_row = route_header + 3
        assert sheet.cell(subtotal_row, 5).value == f"=$D${subtotal_row}+E{route_header + 1}+E{route_header + 2}"
        assert sheet.cell(subtotal_row, 6).value == f"=$D${subtotal_row}+F{route_header + 1}+F{route_header + 2}"
    finally:
        workbook.close()


def test_disabled_lifting_output_removes_the_row_and_compacts_pricing():
    workbook, sheet = _summary_sheet(
        route_rows=[
            {"name": "出厂价", "freight_hkd": "0", "lift_hkd": "0"},
            {"name": "HK 40 柜", "freight_hkd": "1.20", "lift_hkd": "0.50"},
        ],
        lifting_enabled=False,
    )
    try:
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        assert "运费" in values
        route_header = _find_cell(sheet, "运输方案").row
        subtotal_row = route_header + 2
        assert all(
            sheet.cell(row, 2).value != "吊柜费"
            for row in range(route_header, subtotal_row + 1)
        )
        assert sheet.cell(subtotal_row, 5).value == f"=$D${subtotal_row}+E{route_header + 1}"
        assert _find_cell(sheet, "报价（MOQ3K）").row == subtotal_row + 3
    finally:
        workbook.close()


def _assert_printable_grid(sheet, min_row: int, max_row: int, min_col: int, max_col: int):
    for row in range(min_row, max_row + 1):
        for column in range(min_col, max_col + 1):
            border = sheet.cell(row, column).border
            assert border.top.style is not None
            assert border.bottom.style is not None
            assert border.left.style is not None
            assert border.right.style is not None


def test_reference_summary_uses_category_sumifs_and_keeps_sewing_categories_distinct():
    workbook, sheet = _summary_sheet()
    try:
        hair = _find_cell(sheet, "车发")
        cloth_material = _find_cell(sheet, "车衣物料")
        cloth_labor = _find_cell(sheet, "车衣人工")
        assert sheet.cell(hair.row, 2).value == "车发"
        assert sheet.cell(cloth_material.row, 2).value == "车衣"
        assert sheet.cell(cloth_labor.row, 2).value == "车衣"

        summary_header = _find_cell(sheet, "旺季价")
        assert not any(sheet.cell(row, 2).value == "运输方案" for row in range(1, sheet.max_row + 1))
        first_formula = sheet.cell(summary_header.row + 1, 5).value
        for column in range(5, 15):
            header_ref = f"{sheet.cell(summary_header.row, column).column_letter}{summary_header.row}"
            assert sheet.cell(summary_header.row + 1, column).value == first_formula.replace(
                f"E{summary_header.row}", header_ref
            )

        second_header_row = summary_header.row + 3
        second_formula = sheet.cell(second_header_row + 1, 6).value
        for column in range(6, 10):
            header_ref = f"{sheet.cell(second_header_row, column).column_letter}{second_header_row}"
            assert sheet.cell(second_header_row + 1, column).value == second_formula.replace(
                f"F{second_header_row}", header_ref
            )
        assert sheet.cell(second_header_row + 1, 13).value == (
            f"=D{summary_header.row + 1}*$Q$6"
        )
    finally:
        workbook.close()


def test_reference_right_side_tables_have_printable_grid_lines():
    workbook, sheet = _summary_sheet(customer="BUZZ BEE")
    try:
        carton = _find_cell(sheet, "外箱 (inch):")
        _assert_printable_grid(sheet, carton.row, carton.row + 7, 14, 17)

        color_title = _find_cell(sheet, "彩盒价格")
        _assert_printable_grid(sheet, color_title.row, color_title.row + 3, 14, 16)

        test_title = _find_cell(sheet, "测试费用")
        _assert_printable_grid(sheet, test_title.row, test_title.row + 3, 14, 16)
    finally:
        workbook.close()


def test_non_buzzbee_summary_omits_separate_color_box_table_but_keeps_cost_detail():
    workbook, sheet = _summary_sheet(customer="普通客户")
    try:
        values = {
            cell.value
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        }
        assert "彩盒价格" not in values
        assert "报客彩盒" not in values
        assert "报客彩盒FSC" not in values
        assert "彩盒/内卡" in values

        function = next(
            cell
            for row in sheet.iter_rows()
            for cell in row
            if str(cell.value or "").startswith("功能介绍：")
        )
        testing = _find_cell(sheet, "测试费用")
        assert testing.row == function.row + 9
    finally:
        workbook.close()


def test_product_main_image_is_scaled_into_blank_area_above_carton_block():
    workbook, sheet = _summary_sheet(with_image=True)
    try:
        assert len(sheet._images) == 1
        image = sheet._images[0]
        carton = _find_cell(sheet, "外箱 (inch):")
        assert image.anchor._from.col == 13
        assert image.anchor._from.row == 7
        assert image.width <= 340
        assert image.height <= (carton.row - 8) * 20

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        reopened = load_workbook(output, data_only=False)
        try:
            assert len(reopened["报价明细"]._images) == 1
        finally:
            reopened.close()
    finally:
        workbook.close()
