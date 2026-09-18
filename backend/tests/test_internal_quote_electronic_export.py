from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.services.internal_quote import _rr2_cost_summary
from app.services.internal_quote_calculator import calculate_section
from app.services.internal_quote_excel import (
    _build_summary_sheet,
    _justplay_electronic_category,
    _justplay_electronic_source_parts,
)
from app.services.internal_quote_import import parse_internal_quote_workbook
from test_internal_quote_excel_formula_links import _find_row, _section
from test_internal_quote_import import workbook_bytes
from test_internal_quote_purchase_export import arithmetic_value


def _electronic_export_fixture(component_mode=True, inline_battery=False):
    snapshot = {"fx": {"rmb_hkd": ".85", "hkd_usd": "7.8", "rmb_usd": "6.63"},
                "markup": "1.2", "settlement": ".97"}
    imported = parse_internal_quote_workbook(workbook_bytes([
        ["零件名称", "规格", "用量", "单价RMB", "备注"],
        ["IC", "A1", 2, .85],
        ["", "A2", 1, .425],
        ["LED", "红光", 3, .085],
        ["喇叭", "27mm", 1, .255],
        ["PCB", "主板", 1, .85],
        ["电阻", "10K", 1, .085],
        ["邦定成本", .17],
        ["人工成本", .255],
        ["10%利润"],
    ]), "electronic", rmb_hkd=Decimal(".85"))
    payload = imported.payload_fragment
    for row in payload["components"]:
        row["pricing_component_id"] = "a"
    payload["components"].extend([
        {"item": name, "specification": spec, "quantity": 1, "unit_price_rmb": unit,
         "pricing_component_id": "b"}
        for name, spec, unit in [("IC", "B1", .85), ("LED", "蓝光", .17),
                                 ("Speaker", "20mm", .425), ("PCB", "副板", 1.7)]
    ])
    if inline_battery:
        payload["components"].append({"item": "电池", "specification": "旧电子表电池", "quantity": 1,
                                      "unit_price_rmb": .34, "pricing_component_id": "a"})
    electronic = calculate_section("electronic", payload, snapshot, "ELECTRONIC-EXPORT")
    materials = [
        {"item": "AG13电池×3", "category": "auxiliary", "auxiliary_category": "电池",
         "quantity": 3, "unit_price_rmb": .17, "pricing_component_id": "a"},
        {"item": "LR44电池×2", "category": "auxiliary", "auxiliary_category": "电池",
         "quantity": 2, "unit_price_rmb": .255, "pricing_component_id": "b"},
        {"item": "电池片", "category": "hardware", "quantity": 1,
         "unit_price_rmb": .085, "pricing_component_id": "a"},
    ]
    engineering = calculate_section("engineering", {"materials": materials}, snapshot, "ELECTRONIC-EXPORT")
    components = [{"id": "a", "name": "电话", "markup_x": "1.16"},
                  {"id": "b", "name": "镜子", "markup_x": "1.25"}]
    sales = {"pricing_mode": "component", "pricing_components": components} if component_mode else {}
    sections = [_section("electronic", payload=payload, calculation=electronic),
                _section("engineering", payload={"materials": materials}, calculation=engineering),
                _section("sales", payload=sales)]
    for section in sections:
        section.is_required = True
    context = {"electronic_hkd": Decimal(electronic["totals"]["total_hkd"]),
               "auxiliary_hkd": Decimal(engineering["totals"]["auxiliary_hkd"]),
               "hardware_hkd": Decimal(engineering["totals"]["hardware_hkd"])}
    context["factory_price_hkd"] = sum(context.values())
    summary = _rr2_cost_summary(sections, context, snapshot, factory_id="huakang-b")
    workbook = Workbook()
    _build_summary_sheet(workbook, SimpleNamespace(product_name="电子拆分验证", quote_no="JP-ELECTRONIC-QA",
                         customer="JustPlay" if component_mode else "普通客", region_code="mainland", remark="演示文件，不对应正式订单"),
                         sections, snapshot, summary, context, [])
    return workbook, summary, sections, context


def test_imported_parts_split_by_component_without_allocating_overhead_to_ic_led_speaker():
    workbook, summary, sections, context = _electronic_export_fixture()
    sheet = workbook["报价明细"]
    blocks = []
    for component_id, name, raw_costs, battery_label, markup in [
        ("a", "电话", [2.5, .3, .3], "AG13电池×3（3pcs）", 1.16),
        ("b", "镜子", [1, .2, .5], "LR44电池×2（2pcs）", 1.25),
    ]:
        start = _find_row(sheet, 3, f"{name} · 电子")
        cost = next(row for row in range(start + 1, sheet.max_row + 1) if sheet.cell(row, 3).value == "成本金额：")
        labels = [str(sheet.cell(row, 3).value) for row in range(start + 1, cost)]
        for category, expected in zip(("IC", "LED", "喇叭"), raw_costs):
            rows = [row for row in range(start + 1, cost) if _justplay_electronic_category(str(sheet.cell(row, 3).value)) == category]
            assert sum(arithmetic_value(sheet, f"D{row}") for row in rows) == pytest.approx(expected)
            assert all("/$L$4" in sheet.cell(row, 4).value for row in rows)
        if component_id == "a":
            assert any("A1" in label for label in labels) and any("A2" in label for label in labels)
            assert not any("B1" in label for label in labels)
            assert sheet.cell(start + 1, 4).value == "=0.425/$L$4"
            assert sheet.cell(start + 2, 4).value == "=0.85/$L$4*2"
        else:
            assert any("B1" in label for label in labels) and not any("A1" in label for label in labels)
        battery = _find_row(sheet, 3, battery_label)
        assert arithmetic_value(sheet, f"D{battery}") == pytest.approx(.6)
        assert sheet.cell(battery, 2).value == "电池"
        electronic_total = sum(float(entry["amount_hkd"]) for entry in summary["shipping_pricing"]["pricing_entries"]
                               if entry["section"] == "electronic" and entry["pricing_component_id"] == component_id)
        assert arithmetic_value(sheet, f"D{cost}") == pytest.approx(electronic_total + .6)
        pcb = next(row for row in range(start + 1, cost) if sheet.cell(row, 3).value == "PCB其他费用/调整")
        assert sheet.cell(pcb, 4).value == f"=H{start}-SUM(D{start+1}:D{pcb-1})"
        assert sheet.cell(cost + 1, 4).value == markup
        assert sheet.cell(cost + 2, 4).value == "=1-$Q$6"
        primary = _find_row(sheet, 3, f"{name}明细")
        assert not any("电池×" in str(sheet.cell(row, 3).value) for row in range(primary, start))
        assert not any(sheet.cell(row, 3).value == "分配调整" for row in range(primary, start))
        blocks.append((start, cost, electronic_total))
    assert sum(total + .6 for _, _, total in blocks) + .1 == pytest.approx(float(context["factory_price_hkd"]), abs=.0005)
    assert sheet.cell(_find_row(sheet, 3, "电池片（1pcs）"), 2).value == "五金"
    usd_refs = [f"$D${cost+5}" for _, cost, _ in blocks]
    combined = [cell.value for row in sheet for cell in row if cell.data_type == "f" and all(ref in cell.value for ref in usd_refs)]
    assert combined and all(combined[0].count(ref) == 1 for ref in usd_refs)
    for start, cost, _ in blocks:
        for category in ("电子", "电池"):
            headers = [cell for row in sheet for cell in row if cell.value == category and cell.column >= 5]
            assert any(f'SUMIF($B${start+1}:$B${cost-1},{header.coordinate},$D${start+1}:$D${cost-1})'
                       in str(sheet.cell(header.row + 1, header.column).value) for header in headers)
    output = BytesIO()
    workbook.save(output)
    with_formula = load_workbook(BytesIO(output.getvalue()))
    assert with_formula["报价明细"].cell(blocks[0][0] + 1, 4).data_type == "f"
    assert with_formula["报价明细"].cell(blocks[0][0], 1).border.top.style == "medium"
    with_formula.close()
    workbook.close()


def test_old_electronic_source_battery_and_engineering_battery_are_added_only_once():
    workbook, summary, _, _ = _electronic_export_fixture(inline_battery=True)
    sheet = workbook["报价明细"]
    start = _find_row(sheet, 3, "电话 · 电子")
    cost = next(row for row in range(start + 1, sheet.max_row + 1) if sheet.cell(row, 3).value == "成本金额：")
    battery_rows = [row for row in range(start + 1, cost) if sheet.cell(row, 2).value == "电池"]
    assert len(battery_rows) == 2
    assert sum(arithmetic_value(sheet, f"D{row}") for row in battery_rows) == pytest.approx(1)
    total = sum(float(entry["amount_hkd"]) for entry in summary["shipping_pricing"]["pricing_entries"]
                if entry["section"] == "electronic" and entry["pricing_component_id"] == "a")
    assert arithmetic_value(sheet, f"H{start}") == pytest.approx(total - .4)
    assert arithmetic_value(sheet, f"D{cost}") == pytest.approx(total + .6)
    assert "旧电子表电池" in sheet.cell(battery_rows[0], 3).value
    assert "AG13电池×3" in sheet.cell(battery_rows[1], 3).value
    workbook.close()


def test_ordinary_customer_expands_electronic_and_battery_purchase_details():
    workbook, _, _, _ = _electronic_export_fixture(component_mode=False)
    sheet = workbook["报价明细"]
    assert not any(" · 电子" in str(cell.value) for row in sheet for cell in row)
    assert sheet.cell(_find_row(sheet, 3, "A1 IC（2pcs）"), 4).value.startswith("=0.85/$L$4*2")
    assert sheet.cell(_find_row(sheet, 3, "AG13电池×3（3pcs）"), 4).value == "=0.17/$L$4*3"
    workbook.close()


@pytest.mark.parametrize("name,expected", [
    ("IC (240S)", "IC"), ("LED红光", "LED"), ("LEDS", "LED"), ("Speaker 27mm", "喇叭"),
    ("扬声器", "喇叭"), ("AG13电池*3", "电池"), ("BATTERIES", "电池"),
    ("电池片", "PCB"), ("电池盒", "PCB"), ("BATTERY HOLDER", "PCB"),
    ("LED PCB", "PCB"), ("MIC", "PCB"), ("电阻", "PCB"),
])
def test_electronic_export_classification_does_not_mistake_contacts_or_pcb_assemblies_for_batteries(name, expected):
    assert _justplay_electronic_category(name) == expected


@pytest.mark.parametrize("mode", ["detail", "quick", "legacy"])
def test_source_parts_support_quick_and_legacy_hkd_quotes(mode):
    row = {"item": "IC", "specification": "主控", "quantity": 2, "unit_price_hkd": 1,
           "pricing_component_id": "b"}
    payload = {"components": [row]}
    if mode == "detail":
        payload["pricing_currency"] = "RMB"
    if mode == "quick":
        payload = {"quote_mode": "quick", "quick_quotes": [row]}
    calc = calculate_section("electronic", payload, {"fx": {"rmb_hkd": ".85"}}, "SOURCE-PARTS")
    parts = _justplay_electronic_source_parts(_section("electronic", payload=payload, calculation=calc))
    assert len(parts) == 1 and parts[0]["category"] == "IC"
    assert parts[0]["amount_hkd"] == (1 if mode == "quick" else 2)
    assert parts[0]["pricing_component_id"] == "b"
    assert parts[0]["specification"] == "主控"


def test_mismatched_legacy_snapshot_does_not_borrow_payload_category_or_component():
    parts = _justplay_electronic_source_parts(_section("electronic", payload={"components": [{"item": "IC", "pricing_component_id": "a"}]},
        calculation={"line_breakdown": [{"kind": "electronic_component", "item": "LED", "line_hkd": 1, "pricing_component_id": "b"}]}))
    assert parts[0]["category"] == "LED" and parts[0]["pricing_component_id"] == "b"


def test_inconsistent_electronic_total_is_not_silently_exported_as_negative_pcb():
    workbook, summary, sections, context = _electronic_export_fixture()
    workbook.close()
    for entry in summary["shipping_pricing"]["pricing_entries"]:
        if entry["section"] == "electronic":
            entry["amount_hkd"] = 0
    with pytest.raises(ValueError, match="电子总价不足"):
        _build_summary_sheet(Workbook(), SimpleNamespace(product_name="错误报价", quote_no="BAD", customer="JustPlay", region_code="mainland", remark=""),
                             sections, {"fx": {"rmb_hkd": .85, "hkd_usd": 7.8}}, summary, context, [])
