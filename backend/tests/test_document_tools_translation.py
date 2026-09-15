import json
from io import BytesIO

import httpx2 as httpx
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pydantic import SecretStr, ValidationError
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.core.config import settings
from app.schemas.document_tools import CreateJob
from app.services.document_tools import translation_engine as engine
from app.services.document_tools import office_engine as office
from app.services.document_tools.document_ir import Cancelled, ToolError
from test_document_tools_jobs import runtime, upload
from app.services.document_tools import pipeline


def make_pdf(path, texts):
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(600, 800)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 40 750 Td ({text}) Tj ET'.encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
    writer.write(path)


@pytest.fixture
def offline(monkeypatch):
    monkeypatch.setattr(engine.local, "document_translation_status", lambda _: {"available": True})
    def translate(texts, direction, **kwargs):
        return [text.replace("采购订单", "Purchase Order").replace("数量", "Quantity").replace("备注", "Remarks")
                if direction == "zh_to_en" else text.replace("Purchase Order", "采购订单").replace("Quantity", "数量") for text in texts]
    monkeypatch.setattr(engine.local, "translate_texts_locally", translate)
    return translate


def test_word_preserves_source_format_header_and_numbers(tmp_path, offline):
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("采购订单 00123").runs[0].bold = True
    doc.sections[0].header.paragraphs[0].text = "备注"
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "数量 12.50"
    doc.save(path)
    original = path.read_bytes()
    result = engine.convert_translation(path, "word_translate", {}, tmp_path / "work", lambda *a: None, lambda: False)
    output = Document(result.files[0]["path"])
    assert output.paragraphs[0].text == "Purchase Order 00123"
    assert output.paragraphs[0].runs[0].bold
    assert output.sections[0].header.paragraphs[0].text == "Remarks"
    assert output.tables[0].cell(0, 0).text == "Quantity 12.50"
    assert path.read_bytes() == original
    assert result.ir.tables[0].cells[0].raw_text == "数量 12.50"


def test_excel_only_selected_sheet_preserves_formula_and_merge(tmp_path, offline):
    path = tmp_path / "source.xlsx"
    book = Workbook()
    book.active.title = "Selected"
    book.active.append(["采购订单", "00123", 12.5, "=C1*2"])
    book.active.merge_cells("A2:B2")
    book.active["A2"] = "备注"
    book.create_sheet("Untouched")["A1"] = "采购订单"
    book.save(path)
    original = path.read_bytes()
    result = engine.convert_translation(path, "excel_translate", {"sheets": ["Selected"]}, tmp_path / "work", lambda *a: None, lambda: False)
    output = load_workbook(result.files[0]["path"])
    assert output["Selected"]["A1"].value == "Purchase Order"
    assert output["Selected"]["B1"].value == "00123"
    assert output["Selected"]["C1"].value == 12.5
    assert output["Selected"]["D1"].value == "=C1*2"
    assert "A2:B2" in output["Selected"].merged_cells
    assert output["Untouched"]["A1"].value == "采购订单"
    assert path.read_bytes() == original


def test_pdf_translation_reads_selected_page_and_writes_both_formats(tmp_path, monkeypatch, offline):
    path = tmp_path / "source.pdf"
    make_pdf(path, ["Purchase Order 00123", "Quantity 12"])
    monkeypatch.setattr(office, "office_executable", lambda: "test-office")
    def render(source, output, *args):
        assert Document(source).paragraphs[0].text == "数量 12"
        make_pdf(output, ["Rendered translation"])
    monkeypatch.setattr(office, "render_office", render)
    result = engine.convert_translation(path, "pdf_translate", {"translation_direction": "en_to_zh", "page_selection": "2"}, tmp_path / "work", lambda *a: None, lambda: False)
    assert [f["format"] for f in result.files] == ["pdf", "docx"]
    assert len(result.ir.pages) == 1 and result.ir.pages[0].page_index == 1
    assert result.ir.blocks[0].original_text == "Quantity 12"


@pytest.mark.parametrize("protocol", ["dashscope", "openai"])
def test_online_contract_and_glossary(monkeypatch, protocol):
    monkeypatch.setattr(settings, "document_tools_ai_mode", "auto")
    monkeypatch.setattr(settings, "document_tools_qwen_api_key", SecretStr("test-key"))
    monkeypatch.setattr(settings, "document_tools_qwen_base_url", "https://example.test/api/v1" if protocol == "dashscope" else "https://example.test/v1")
    monkeypatch.setattr(settings, "document_tools_qwen_protocol", protocol)
    def handle(request):
        body = json.loads(request.content)
        messages = body["input"]["messages"] if protocol == "dashscope" else body["messages"]
        assert messages[0]["role"] == "system"
        assert "terminology" in json.dumps(messages)
        assert request.headers["Authorization"] == "Bearer test-key"
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {"content": '{"translations":["Quantity 12"]}'}}]})
    with httpx.Client(transport=httpx.MockTransport(handle)) as client:
        assert engine.online_batch(["数量 12"], "zh_to_en", "数量=Quantity", lambda: False, client) == ["Quantity 12"]


@pytest.mark.parametrize("response", [{"translations": []}, {"translations": [""]}, {"translations": "wrong"}])
def test_invalid_ai_output_is_not_published(monkeypatch, response):
    monkeypatch.setattr(engine, "online_configured", lambda: True)
    monkeypatch.setattr(settings, "document_tools_qwen_base_url", "https://example.test/v1")
    with httpx.Client(transport=httpx.MockTransport(lambda _: httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(response)}}]}))) as client:
        with pytest.raises(ToolError, match="不完整"):
            engine.online_batch(["数量"], "zh_to_en", "", lambda: False, client)


def test_number_changes_and_cancellation_rejected(monkeypatch, offline):
    monkeypatch.setattr(engine, "online_batch", lambda *a: ["Quantity 124"])
    with pytest.raises(ToolError, match="数字"):
        engine.make_translator({"translation_engine": "online"}, lambda *a: None, lambda: False)(["数量 123"], "zh_to_en")
    with pytest.raises(Cancelled):
        engine.make_translator({}, lambda *a: None, lambda: True)(["数量"], "zh_to_en")


def test_translation_schema_rejects_unrelated_options():
    with pytest.raises(ValidationError):
        CreateJob(source_id="s", operation="excel_translate", client_request_id="c", options={"formula_mode": "display"})


@pytest.mark.parametrize("source, translated", [("温度 -20 C", "Temperature 20 C"), ("税率 12%", "Rate 12"), ("电流 1e-3 A", "Current 1e3 A"), ("型号 AB123", "Model CD123")])
def test_numeric_semantics_and_model_codes_are_protected(monkeypatch, source, translated):
    monkeypatch.setattr(engine, "online_batch", lambda *a: [translated])
    with pytest.raises(ToolError):
        engine.make_translator({"translation_engine": "online"}, lambda *a: None, lambda: False)([source], "zh_to_en")


def test_job_translation_download_owner_scope_and_revision_guard(runtime, offline):
    client, sessions, current = runtime
    doc = Document()
    doc.add_paragraph("采购订单 00123")
    content = BytesIO()
    doc.save(content)
    data = upload(client, data=content.getvalue())
    pipeline.run_one(sessions)
    created = client.post("/api/tools/jobs", json={"source_id": data["source_id"], "operation": "word_translate", "options": {}, "client_request_id": "translate-test"})
    assert created.status_code == 202, created.text
    pipeline.run_one(sessions)
    job_id = created.json()["job_id"]
    result = client.get(f"/api/tools/jobs/{job_id}").json()
    assert result["execution_status"] == "succeeded", result
    artifact = next(a for a in result["artifacts"] if a["role"] == "result")
    assert artifact["filename"] == "table_中译英.docx"
    downloaded = client.get(f'/api/tools/artifacts/{artifact["id"]}/download')
    assert Document(BytesIO(downloaded.content)).paragraphs[0].text == "Purchase Order 00123"
    assert client.post(f"/api/tools/jobs/{job_id}/revise", json={"base_revision": 1, "options": {"translation_direction": "en_to_zh"}}).status_code == 422
    current.id = "bob"
    assert client.get(f'/api/tools/artifacts/{artifact["id"]}/download').status_code == 404


def test_unconfigured_online_rejected_before_enqueue(runtime):
    client, sessions, _ = runtime
    data = upload(client)
    pipeline.run_one(sessions)
    response = client.post("/api/tools/jobs", json={"source_id": data["source_id"], "operation": "word_translate", "options": {"translation_engine": "online"}, "client_request_id": "unconfigured"})
    assert response.status_code == 422
    assert "TRANSLATION_NOT_CONFIGURED" in response.text
