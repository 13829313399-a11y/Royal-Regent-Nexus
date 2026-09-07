from io import BytesIO

import openpyxl
import pytest

from app.services import customer_order_unified as service
from app.services import customer_order_regional as mapping


def template():
    w = openpyxl.Workbook()
    w.remove(w.active)
    for name, marker in ((service.ITEM_SHEET, 7), (service.ORDER_SHEET, 5), (service.REVIEW_SHEET, 6)):
        s = w.create_sheet(name)
        for c, v in enumerate(service.ITEM_HEADERS if name == service.ITEM_SHEET else service.ORDER_HEADERS, 1):
            s.cell(3, c, v)
        s.cell(marker, 1, '取消单')
        s.cell(marker+2, 1, '已走货订单')
        s.cell(marker+3, 1, '原历史备注')
        s.cell(4, 3, '原生产号')
    item = w[service.ITEM_SHEET]
    item['F10'], item['H10'], item['K10'], item['L10'] = 'JAZ1', 'ITEM1', 492, 12
    item['K10'].number_format = item['L10'].number_format = 'yyyy/mm/dd'
    b=BytesIO()
    w.save(b)
    return b.getvalue()


def test_history_po_fallback_and_numeric_date_format():
    history = service.read_unified_history(template(), factory_id='huakang-d', customer_code='jazwares')
    assert len(history) == 1
    assert (history[0].reference_no, history[0].quantity, history[0].units_per_carton) == ('JAZ1', '492', '12')
    w=openpyxl.load_workbook(BytesIO(template()))
    w[service.ITEM_SHEET]['L10']='12/48'
    b=BytesIO(); w.save(b)
    history=service.read_unified_history(b.getvalue(),factory_id='huadeng',customer_code='simba')
    assert history[0].units_per_carton == '48'


@pytest.mark.parametrize('factory,customer',[('huadeng','spin-master'),('huakang-d','jp')])
def test_removed_customer_entry_points_reject_new_imports(factory,customer):
    from fastapi import HTTPException
    from app.api import customer_order as api
    with pytest.raises(HTTPException):
        api._ensure_customer_factory(customer,factory)
    assert not service._factory_allowed(customer,factory)
    assert not mapping.enabled(factory,customer)
    assert service._factory_allowed('spin','huadeng')
    api._ensure_customer_factory('spin','huadeng')
    api._ensure_customer_factory('jp','huakang-c')


@pytest.mark.parametrize('customer,record,reference,product', [
    ('jakks', {'confirmation_no':'5460525','contract_no':'808464','special_note':'请在周一送货'}, '5460525','SKU'),
    ('simba', {'contract_no':'700147991/1700','contract_line':'1700','certificate':'NO'}, '700147991/1700','SKU'),
    ('spin', {'contract_no':'4500/10','unified_reference':'37594/1099/6077/3837-11','material_group':'37594'}, '37594/1099/6077/3837-11','37594'),
    ('spin-master', {'contract_no':'4500/10','unified_reference':'37594/1099/6077/3837-11','material_group':'37594'}, '37594/1099/6077/3837-11','37594'),
    ('casdon', {'contract_no':'0008785','po_no':'P1453'}, '0008785','SKU'),
])
def test_customer_reference_and_semantic_fields(customer,record,reference,product):
    row=mapping.map_huadeng({'product_no':'SKU','standard':'旧展示字符串','packaging':'旧组合备注'},record,customer)
    assert (row['reference_no'],row['product_no']) == (reference,product)
    assert row['standard'] == row['packaging'] == ''
    if customer == 'casdon':
        assert row['po_no'] == '0008785'
        assert row['end_customer_po'] == 'P1453'


def test_casdon_suffix_cannot_cross_customer_or_version():
    h=service.UnifiedHistoryRow(row=4,reference_no='1',customer_name='Target',product_no='11050.TAR001（AW26）')
    row={'reference_no':'2','customer_name':'KB SALES','product_no':'11050(AW26)'}
    mapping.reconcile_product(row,[h],'huadeng','casdon')
    assert row['product_no']=='11050(AW26)'
    row['customer_name']='Target'
    mapping.reconcile_product(row,[h],'huadeng','casdon')
    assert row['product_no']==h.product_no


def test_jakks_missing_confirmation_reconciles_existing_order():
    h=service.UnifiedHistoryRow(row=177,contract_no='808464',reference_no='5460525',product_no='39664-2L-RF1',quantity='300')
    row=mapping.map_huadeng({'id':'j1','contract_no':'808464','product_no':h.product_no,'quantity':'300','issues':[]},{'contract_no':'808464'},'jakks')
    mapping.reconcile_product(row,[h],'huadeng','jakks')
    service._enrich_row(row,[h],'Jakks')
    assert row['reference_no']=='5460525'
    assert any(i['code']=='duplicate_existing_order' for i in row['issues'])
    h.contract_no=''
    row=mapping.map_huadeng({'id':'j2','product_no':h.product_no,'issues':[]},{},'jakks')
    mapping.reconcile_product(row,[h],'huadeng','jakks')
    service._enrich_row(row,[h],'Jakks')
    assert any(i['code']=='missing_contract_no' for i in row['issues'])


def test_jakks_kilogram_contract_is_not_silently_treated_as_pieces():
    row=mapping.map_huadeng({'id':'kg','quantity':'528.32','issues':[]},{'contract_no':'HD20260604','unit':'KGM'},'jakks')
    assert any(i['code']=='unsupported_quantity_unit' for i in row['issues'])


def test_jakks_explicit_po_standard_and_packaging_override_history(monkeypatch):
    from app.services import customer_order_unified as current_service
    from app.services.huadeng_order_legacy import jakks_po_parser
    text = '''CONTRACT
JAKKS PACIFIC (H.K.) LIMITED
Contract #: 812136
Order Date: 24-Jun-2026
Customer PO: 0107993764
Confirmation No: 66959 GERMANY
Ultimate Consignee Name: JAKKS EUROPE BV
ITEM NUMBER DESCRIPTION QUANTITY PRICE EXTENDED PRICE
11180J-V2 GIFT COLLECTION-GEN- 24 1.1842 28.42 USD
PQ 2 CA
Customer Item No.: 11180J-V2
Outer Pack: 12
Country of Origin: CHINA Line#: 1.000 Request Date: 15-Sep-2026
NOTES: WINDOW BOX
PRODUCT MEETS EU STANDARD.
TOTAL USD 28.42
'''
    monkeypatch.setattr(jakks_po_parser, 'extract_text', lambda _: (text, False))
    w=openpyxl.load_workbook(BytesIO(template()))
    s=w[current_service.ITEM_SHEET]
    for col,value in {4:'OLD',5:'OLD',7:'JAKKS EUROPE BV GERMANY',8:'11180J-V2',9:'中文品名',11:24,12:12,14:'AS/NZS ISO8124',20:'PDQ盒'}.items():
        s.cell(4,col,value)
    b=BytesIO();w.save(b)
    output,_,preview=current_service.export_unified_customer_schedule(factory_id='huadeng',customer_code='jakks',received_date='2026-09-07',po_files=[('formal.pdf',b'pdf')],schedule_file_name='latest.xlsx',schedule_content=b.getvalue())
    row=preview['rows'][0]
    assert row['product_name_en']=='GIFT COLLECTION-GEN-PQ'
    assert row['carton_count']=='2'
    assert row['standard']=='EU STANDARD'
    assert row['packaging']=='WINDOW BOX'
    assert row['po_no']=='0107993764'
    saved=openpyxl.load_workbook(BytesIO(output))[current_service.ITEM_SHEET]
    assert (saved['N5'].value,saved['T5'].value)==('EU STANDARD','WINDOW BOX')
    assert (saved['N4'].value,saved['T4'].value)==('AS/NZS ISO8124','PDQ盒')


def test_maxx_absent_packing_is_blank_instead_of_zero():
    row=mapping.map_huakang_d({'quantity':'6000','units_per_carton':'0','carton_count':'0'}, {'po_number':'PO1'}, {}, 'maxx')
    assert row['units_per_carton']==row['carton_count']==''


def test_new_summary_rows_link_exact_item_despite_broken_or_duplicate_lookup():
    w=openpyxl.load_workbook(BytesIO(template()))
    s=w[service.ORDER_SHEET]
    s['H4']='=COUNTIFS(ITEM表!E:E,D4)'
    s['H5']='=IFERROR(INDEX(ITEM表!#REF!,1),"")'
    s['N5']=None
    s['P5']='=N5/7.8'
    service._write_summary_row(s,5,{'reference_no':'DUP','product_no':'SKU'},item_row=20,item_sheet=w[service.ITEM_SHEET])
    assert s['H4'].value=='=COUNTIFS(ITEM表!E:E,D4)'
    assert s['H5'].value=='=IF(LEN(\'ITEM表\'!I20)=0,"",\'ITEM表\'!I20)'
    assert s['J5'].value=='=IF(LEN(\'ITEM表\'!K20)=0,"",\'ITEM表\'!K20)'
    assert s['N5'].value is None
    assert s['P5'].value=='=N5/7.8'


def test_generated_wrapped_text_height_keeps_original_column_dimensions():
    from openpyxl.styles import Alignment
    w=openpyxl.Workbook();s=w.active
    s.column_dimensions.group('A','C')
    s.column_dimensions['A'].width=25
    s['B5']='客户名字很长，需要保留完整显示，测试自动扩展新增行高度。'
    s['B5'].alignment=Alignment(wrap_text=True)
    dimensions=list(s.column_dimensions)
    service._fit_generated_text(s,5)
    assert s.row_dimensions[5].height>30
    assert list(s.column_dimensions)==dimensions
    assert s.column_dimensions['A'].width==25


def test_inserted_numeric_cells_do_not_inherit_historical_date_display():
    from app.services.huakang_a_order_legacy.unified_schedule import output_slots
    w=openpyxl.load_workbook(BytesIO(template()))
    s=w[service.ITEM_SHEET]
    s['K4'],s['L4']=492,12
    s['K4'].number_format=s['L4'].number_format=s['M4'].number_format='yyyy/mm/dd'
    slots=output_slots(w,6)
    for r in slots[service.ITEM_SHEET]:
        service._write_item_row(s,r,{'quantity':'15000','units_per_carton':'6','_regional_customer':'jazwares'})
    b=BytesIO();w.save(b)
    saved=openpyxl.load_workbook(BytesIO(b.getvalue()))[service.ITEM_SHEET]
    assert saved['K4'].is_date
    assert all(saved.cell(r,11).value==15000 and saved.cell(r,12).value==6 for r in slots[service.ITEM_SHEET])


def test_casdon_extras_follow_headers_without_overwriting_manual_columns():
    w=openpyxl.load_workbook(BytesIO(template())); s=w[service.ITEM_SHEET]
    s['AE3'],s['AH3'],s['AL3']='客PO','品牌','验货报告编码'
    s['AL5']='人工验货记录'
    mapping.write_extras(s,5,{'end_customer_po':'P1453','brand':'Build-A-Bear'},'huadeng','casdon')
    assert (s['AE5'].value,s['AH5'].value,s['AL5'].value)==('P1453','Build-A-Bear','人工验货记录')


@pytest.mark.parametrize('factory,customer',[('huadeng','jakks'),('huakang-d','jazwares')])
def test_export_independent_sheet_rows_and_history_conflicts(monkeypatch,factory,customer):
    def parser(*args):
        return [{'id':'new-1','reference_no':'NEW','contract_no':'NEW','product_no':'ITEM1','quantity':'492','issues':[], '_regional_customer':customer}], [], 'TEST'
    monkeypatch.setattr(service,'_create_customer_rows',parser)
    args=dict(factory_id=factory,customer_code=customer,received_date='2026-09-07',po_files=[('new.pdf',b'po')],schedule_file_name='new.xlsx',schedule_content=template())
    output,_,preview=service.export_unified_customer_schedule(**args)
    assert preview['target_template']==mapping.TARGET_TEMPLATE
    w=openpyxl.load_workbook(BytesIO(output))
    for name in service.SHEETS:
        assert w[name]['C4'].value=='原生产号'
        assert service._marker_row(w[name])>5
        assert any(c.value=='原历史备注' for c in w[name]['A'])
    assert w[service.ITEM_SHEET]['E5'].value=='NEW'
    assert w[service.ORDER_SHEET]['D5'].value=='NEW'
    assert w[service.ITEM_SHEET]['K5'].number_format=='General'
    assert w[service.ITEM_SHEET]['K10'].is_date
    row={'id':'x','reference_no':'JAZ1','product_no':'ITEM1','quantity':'100','issues':[]}
    service._enrich_row(row,service.read_unified_history(template(),factory_id=factory,customer_code=customer),'Customer')
    assert any(i['code']=='existing_quantity_conflict' for i in row['issues'])
