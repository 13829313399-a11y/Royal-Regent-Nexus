import importlib
from copy import deepcopy
import io
import json
import zipfile
from contextlib import nullcontext
from unittest.mock import MagicMock

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from app.services import fabric_legacy_stock as service

HEADERS = ['供应商', '布料', '入库日期', '位置', '新物料编号', '物料编号', '物料名称规格', '单位',
           '客名', '合同', '放产编号', '货号', '备注单价', '单价', '上月结存', '本月结存', '是否暂停', '仓库分析', '2026报减值']


def line(**overrides):
    return {'供应商': '测试纺织厂', '布料': '布料', '入库日期': '2026-09-01', '位置': '布料一仓／A01',
            '新物料编号': '00012', '物料编号': 'OLD-12', '物料名称规格': '白布58寸', '单位': '码',
            '客名': '测试客户', '合同': '00009', '放产编号': '01-02-0003', '货号': '0007',
            '备注单价': '按原单位', '单价': '3.25', '上月结存': '99', '本月结存': '12.345678',
            '是否暂停': '否', '仓库分析': '待核对', '2026报减值': '0', **overrides}


def book(rows=None, headers=HEADERS, offset=2):
    workbook = Workbook(); sheet = workbook.active; sheet.title = '物料'
    for _ in range(offset):
        sheet.append(['旧库存表说明'])
    sheet.append(headers)
    for record in rows if rows is not None else [line()]:
        sheet.append([record.get(header) for header in headers])
    output = io.BytesIO(); workbook.save(output)
    return output.getvalue()


def masters():
    return [dict(id='M1', kind='MATERIAL', code='00012', name='白布58寸', status='ACTIVE', data={'old_code': 'OLD-12'}),
            dict(id='S1', kind='SUPPLIER', code='SUP-1', name='测试纺织厂', status='ACTIVE', data={})]


def locations():
    return [dict(id='L1', label='布料一仓／A01', warehouse='布料一仓', code='A01', status='ACTIVE')]


def inspect(content=None, **kwargs):
    return service.inspect_rows(service.parse('legacy.xlsx', content or book()), kwargs.get('masters', masters()), kwargs.get('locations', locations()))


def test_header_offsets_reordered_columns_and_reference_mapping():
    parsed = service.parse('legacy.xlsx', book(headers=HEADERS[::-1], offset=5))
    assert parsed['header_row'] == 6 and parsed['quantity_column'] == '本月结存'
    assert parsed['field_mapping']['current_quantity'] == {'header': '本月结存', 'column': 'D'}
    row = parsed['rows'][0]
    assert row['row_number'] == 7
    assert row['values']['contract'] == '00009' and row['values']['production_no'] == '01-02-0003'
    assert row['values']['price_note'] == '按原单位' and row['values']['unit_price'] == '3.25'
    assert inspect(book(headers=HEADERS[::-1], offset=5))['summary']['valid_rows'] == 1


def test_codes_keep_leading_zeros_in_text_and_explicit_numeric_formats():
    workbook = Workbook(); sheet = workbook.active; sheet.title = '物料'; sheet.append(HEADERS)
    sheet.append([line().get(key) for key in HEADERS])
    sheet['E2'] = 12; sheet['E2'].number_format = '00000'
    output = io.BytesIO(); workbook.save(output)
    result = inspect(output.getvalue())
    assert result['summary']['valid_rows'] == 1
    parsed = service.parse('legacy.xlsx', output.getvalue())['rows'][0]['values']
    assert parsed['new_item_code'] == '00012' and parsed['style_no'] == '0007'
    assert parsed['production_no'] == '01-02-0003'


def test_zero_missing_negative_nonnumeric_and_unit_totals_are_distinct():
    result = inspect(book([line(**{'本月结存': value, '合同': str(index)}) for index, value in enumerate([0, None, -1, '不是数字', 'NaN', 'Infinity', True])]
                          + [line(**{'本月结存': '0.1', '单位': '米'}), line(**{'本月结存': '0.2', '单位': '米'}), line(**{'本月结存': '2', '单位': ''})]))
    assert result['summary'] == dict(total_rows=10, valid_rows=3, zero_rows=1, missing_qty=1, invalid_qty=5, missing_unit=1,
                                     duplicate_rows=0, duplicate_groups=0, conversion_only_rows=0)
    assert result['totals'] == {'码': '0', '米': '0.3'}
    assert result['problems'][0]['issues'] == ['MISSING_QTY']
    assert result['problems'][0]['raw']['previous_quantity'] == '99'


def test_prior_month_only_used_if_current_column_is_absent():
    result = inspect(book(headers=[key for key in HEADERS if key != '本月结存']))
    assert result['quantity_column'] == '上月结存' and result['totals'] == {'码': '99'}


def test_formula_uses_saved_cache_and_never_evaluates_or_fills_missing_cache():
    content = book([line(**{'本月结存': '=Z4+AB4-AC4'}), line(**{'本月结存': '=1+1', '合同': 'different'})])
    source = zipfile.ZipFile(io.BytesIO(content)); output = io.BytesIO()
    with source, zipfile.ZipFile(output, 'w') as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == 'xl/worksheets/sheet1.xml':
                data = data.replace(b'<f>Z4+AB4-AC4</f><v></v>', b'<f>Z4+AB4-AC4</f><v>0</v>')
            target.writestr(info, data)
    result = inspect(output.getvalue())
    assert result['summary']['zero_rows'] == 1 and result['summary']['missing_qty'] == 1
    assert result['totals'] == {'码': '0'}
    assert result['problems'][0]['raw']['current_quantity'] == '=1+1'


def test_unmatched_ambiguous_and_inactive_matches_do_not_guess():
    missing = inspect(book([line(**{'新物料编号': 'missing', '物料编号': '', '供应商': '', '位置': 'A01'})]))
    assert set(missing['problems'][0]['issues']) == {'UNMATCHED_MATERIAL', 'UNMATCHED_SUPPLIER', 'UNMATCHED_LOCATION'}
    records = masters() + [dict(id='M2', kind='MATERIAL', code='OTHER', name='其他布', status='ACTIVE', data={'old_code': 'OLD-12'}),
                           dict(id='S2', kind='SUPPLIER', code='SUP-2', name='测试纺织厂', status='ACTIVE', data={})]
    places = locations() + [{**locations()[0], 'id': 'L2'}]
    ambiguous = inspect(masters=records, locations=places)
    assert set(ambiguous['problems'][0]['issues']) == {'AMBIGUOUS_MATERIAL', 'AMBIGUOUS_SUPPLIER', 'AMBIGUOUS_LOCATION'}
    records = [{**row, 'status': 'DRAFT'} for row in masters()]
    inactive = inspect(masters=records, locations=[{**locations()[0], 'status': 'INACTIVE'}])
    assert set(inactive['problems'][0]['issues']) == {'MATERIAL_NOT_ACTIVE', 'SUPPLIER_NOT_ACTIVE', 'UNMATCHED_LOCATION'}
    alias = inspect(book([line(**{'新物料编号': '', '供应商': 'SUP-1', '位置': '布料一仓/A01'})]))
    assert alias['summary']['valid_rows'] == 1


@pytest.mark.parametrize('row_count', [207, 503])
def test_duplicate_rows_and_problem_cap_keep_full_counts_and_source_totals(row_count):
    result = inspect(book([line()] * row_count), masters=[], locations=[])
    assert result['summary']['duplicate_rows'] == row_count - 1
    assert result['problem_count'] == row_count and len(result['problems']) == min(row_count, 500)
    assert result['problems_truncated'] == (row_count > 500)
    assert result['matching']['material'] == dict(matched=0, unmatched=row_count, ambiguous=0)
    assert result['problems'][1]['duplicate_of'] == 4
    assert result['totals'] == {'码': '2555.555346' if row_count == 207 else '6209.876034'}
    assert result['unique_totals'] == {'码': '12.345678'} and result['matched_unique_totals'] == {'码': '0'}
    assert result['summary']['duplicate_groups'] == 1 and not result['duplicate_groups_truncated']
    group = result['duplicate_groups'][0]
    assert group['count'] == row_count and not group['rows_truncated']
    assert [row['row_number'] for row in group['rows']] == list(range(4, row_count + 4))
    assert group['rows'][0]['duplicate_of'] is None
    assert all(row['duplicate_of'] == 4 for row in group['rows'][1:])
    assert group['rows'][0]['supplier'] == '测试纺织厂'


def test_unit_aliases_preserve_original_rows_quantities_and_duplicate_identity():
    units = ['Pcs', 'PCS', 'pcs', 'KG', 'kg', 'Kg', '千克', '公斤', '码', '卷', '个', '套', '对', '条', '只', '片', '支']
    parsed = service.parse('legacy.xlsx', book([line(**{'单位': unit, '本月结存': '0.1'}) for unit in units]))
    before = deepcopy(parsed)
    result = service.inspect_rows(parsed, masters(), locations())
    assert parsed == before
    assert result['totals'] == {'PCS': '0.3', 'KG': '0.5', **{unit: '0.1' for unit in units[8:]}}
    assert result['unique_totals'] == result['matched_unique_totals'] == result['totals']
    assert result['summary']['duplicate_rows'] == 0 and result['duplicate_groups'] == []
    assert [row['raw_unit'] for row in result['rows']] == units
    assert [row['unit'] for row in result['rows']] == ['PCS'] * 3 + ['KG'] * 5 + units[8:]
    assert all(row['quantity'] == '0.1' and row['unit_basis'] and not row['conversion_only'] for row in result['rows'])
    assert result['unit_aliases'][0] == dict(raw='Pcs', canonical='PCS', rows=1, conversion_only=False)


def test_placeholder_units_are_missing_and_excluded_from_all_totals():
    units = ['0', '-', '—', '/', '无', '', '  ']
    result = inspect(book([line(**{'单位': unit}) for unit in units]))
    assert result['summary']['missing_unit'] == len(units)
    assert all('MISSING_UNIT' in row['issues'] and row['unit'] == '' for row in result['problems'])
    assert result['totals'] == result['unique_totals'] == result['matched_unique_totals'] == result['conversion_totals'] == {}
    assert all(alias['canonical'] == '' for alias in result['unit_aliases'])
    assert sum(alias['rows'] for alias in result['unit_aliases']) == len(units)


def test_conversion_only_units_are_separate_even_for_duplicate_and_invalid_rows():
    result = inspect(book([line(**{'单位': '千克', '本月结存': '3'}),
                           line(**{'单位': '千克(转算专用)', '本月结存': '5'}) ,
                           line(**{'单位': '千克(转算专用)', '本月结存': '5'}),
                           line(**{'单位': 'KG（转算专用）', '本月结存': '-1'})]))
    assert result['totals'] == result['unique_totals'] == result['matched_unique_totals'] == {'KG': '3'}
    assert result['conversion_totals'] == {'KG（转算专用）': '10'}
    assert result['summary']['conversion_only_rows'] == 3 and result['summary']['duplicate_rows'] == 1
    assert result['unit_aliases'][1] == dict(raw='千克(转算专用)', canonical='KG（转算专用）', rows=2, conversion_only=True)
    assert result['rows'][1]['quantity'] == '5' and result['rows'][1]['conversion_only']
    assert result['duplicate_groups'][0]['rows'][0]['unit'] == 'KG（转算专用）'


def test_source_unique_and_matched_unique_totals_have_explicit_distinct_scopes():
    result = inspect(book([line(**{'单位': 'PCS', '本月结存': '10'})] * 2 + [
        line(**{'单位': 'PCS', '本月结存': '5', '新物料编号': 'UNKNOWN', '物料编号': ''}),
        line(**{'单位': 'KG', '本月结存': '2'}), line(**{'单位': '0', '本月结存': '9'}),
        line(**{'单位': 'PCS', '本月结存': '-1'}), line(**{'单位': '千克(转算专用)', '本月结存': '100'})]))
    assert result['totals'] == {'PCS': '25', 'KG': '2'}
    assert result['unique_totals'] == {'PCS': '15', 'KG': '2'}
    assert result['matched_unique_totals'] == {'PCS': '10', 'KG': '2'}
    assert result['conversion_totals'] == {'KG（转算专用）': '100'}
    assert '不能直接作为系统库存' in result['totals_basis']
    inactive = inspect(masters=[{**record, 'status': 'INACTIVE'} for record in masters()])
    assert inactive['matching']['material']['matched'] == 1
    assert inactive['matched_unique_totals'] == {'码': '0'}


@pytest.mark.parametrize('group_count', [200, 201])
def test_duplicate_group_limit_keeps_complete_statistics(group_count):
    result = inspect(book([line(**{'合同': str(index)}) for index in range(group_count) for _ in range(2)]))
    assert result['summary']['duplicate_groups'] == result['summary']['duplicate_rows'] == group_count
    assert len(result['duplicate_groups']) == min(group_count, 200)
    assert result['duplicate_groups_truncated'] == (group_count > 200)
    assert all(group['count'] == 2 and len(group['rows']) == 2 for group in result['duplicate_groups'])


@pytest.mark.parametrize('row_count', [1000, 1001])
def test_duplicate_row_limit_includes_first_row_and_marks_partial_groups(row_count):
    result = inspect(book([line()] * row_count))
    assert result['summary']['duplicate_rows'] == row_count - 1
    assert result['duplicate_groups'][0]['count'] == row_count
    assert len(result['duplicate_groups'][0]['rows']) == min(row_count, 1000)
    assert result['duplicate_groups'][0]['rows_truncated'] == result['duplicate_groups_truncated'] == (row_count > 1000)
    assert len(result['rows']) == row_count and result['unique_totals'] == {'码': '12.345678'}


def test_duplicate_row_budget_is_shared_between_groups():
    result = inspect(book([line(**{'合同': str(index)}) for index in range(2) for _ in range(600)]))
    assert [len(group['rows']) for group in result['duplicate_groups']] == [600, 400]
    assert [group['count'] for group in result['duplicate_groups']] == [600, 600]
    assert result['duplicate_groups_truncated'] and result['summary']['duplicate_rows'] == 1198


@pytest.mark.parametrize('filename,content,status', [('legacy.xls', b'x', 422), ('legacy.xlsx', b'bad', 422), ('legacy.xlsx', b'x' * (service.MAX_BYTES + 1), 413)], ids=['extension', 'corrupt', 'oversized'])
def test_invalid_or_oversized_files(filename, content, status):
    with pytest.raises(HTTPException) as caught:
        service.parse(filename, content)
    assert caught.value.status_code == status


def test_row_limit_missing_sheet_and_duplicate_headers():
    with pytest.raises(HTTPException, match='5000'):
        service.parse('legacy.xlsx', book([line()] * 5001, offset=0))
    with pytest.raises(HTTPException, match='表头重复'):
        service.parse('legacy.xlsx', book(headers=HEADERS + ['单位']))
    workbook = Workbook(); output = io.BytesIO(); workbook.save(output)
    with pytest.raises(HTTPException, match='物料'):
        service.parse('legacy.xlsx', output.getvalue())


def test_service_only_reads_catalogs_without_posting_or_database_writes(monkeypatch):
    db = MagicMock(); db.no_autoflush = nullcontext()
    monkeypatch.setattr(service.fabric_master, 'ensure_schema', lambda _: None)
    monkeypatch.setattr(service.fabric_master, 'records', lambda _: masters())
    monkeypatch.setattr(service.fabric_master, 'record_data', lambda record: record)
    monkeypatch.setattr(service.warehouse_locations, 'catalog', lambda *_: locations())
    ops = importlib.import_module('app.services.warehouse_operations')
    setup = importlib.import_module('app.services.fabric_inventory_setup')
    post = MagicMock(side_effect=AssertionError('posting is forbidden')); apply = MagicMock(side_effect=AssertionError('opening is forbidden'))
    monkeypatch.setattr(ops, 'post', post); monkeypatch.setattr(setup, 'apply_opening', apply)
    result = service.preview(db, 'legacy.xlsx', book())
    assert result['read_only'] is True and result['posted'] is False
    post.assert_not_called(); apply.assert_not_called()
    for method in ('add', 'add_all', 'delete', 'flush', 'commit', 'execute'):
        getattr(db, method).assert_not_called()


def test_endpoint_read_permission_factory_boundary_and_no_sql_writes(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from sqlalchemy import create_engine, event
    from sqlalchemy.orm import Session
    from sqlalchemy.pool import StaticPool
    from app.models import fabric_master as master, fabric_procurement as procurement

    # Exercise the actual router and SQL readers without starting unrelated app modules.
    engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    models = (procurement.FabricProcurementState, procurement.FabricProcurementLine,
              procurement.FabricProcurementImport, procurement.FabricProcurementEvidence,
              master.FabricMasterRecord, master.FabricMasterChange, master.FabricChaseResolution)
    for model in models:
        model.__table__.create(engine)
    with Session(engine) as db:
        for record in masters() + [dict(id='L1', kind='LOCATION', code='A01', name='A01', status='ACTIVE', data={'warehouse': '布料一仓'})]:
            db.add(master.FabricMasterRecord(**{key: record[key] for key in ('id', 'kind', 'code', 'name', 'status')},
                                            factory_id='huakang-c', revision=1, data_json=json.dumps(record['data']), updated_at='2026-10-23'))
        db.commit()
    api = importlib.import_module('app.api.warehouse_operations')
    app = FastAPI(); app.include_router(api.router)
    def database():
        with Session(engine) as db:
            yield db
    app.dependency_overrides[api.get_db] = database
    app.dependency_overrides[api.get_current_user] = lambda: object()
    monkeypatch.setattr(api, 'can', lambda actor, permission, *args: permission == 'fabric_operations:read')
    statements = []
    def capture(conn, cursor, statement, parameters, context, executemany):
        statements.append(statement.strip().split()[0].upper())
    event.listen(engine, 'before_cursor_execute', capture)
    try:
        with TestClient(app) as client:
            response = client.post('/api/warehouse-operations/fabric/legacy-stock/preview', data={'factory_id': 'huakang-c'}, files={'file': ('legacy.xlsx', book())})
            assert response.status_code == 200, response.text
            assert response.json()['read_only'] and not response.json()['posted']
            assert response.json()['summary']['valid_rows'] == 1
            assert response.headers['cache-control'] == 'no-store'
            assert client.post('/api/warehouse-operations/fabric/legacy-stock/preview', data={'factory_id': 'huakang-b'}, files={'file': ('legacy.xlsx', book())}).status_code in (403, 422)
            assert client.post('/api/warehouse-operations/fabric/legacy-stock/preview', data={'factory_id': 'huakang-c'}, files={'file': ('legacy.xlsx', b'x' * (service.MAX_BYTES + 1))}).status_code == 413
            monkeypatch.setattr(api, 'can', lambda *_: False)
            assert client.post('/api/warehouse-operations/fabric/legacy-stock/preview', data={'factory_id': 'huakang-c'}, files={'file': ('legacy.xlsx', book())}).status_code == 403
        assert statements and not set(statements) & {'INSERT', 'UPDATE', 'DELETE', 'CREATE', 'ALTER', 'DROP', 'REPLACE'}
    finally:
        event.remove(engine, 'before_cursor_execute', capture)
        engine.dispose()
