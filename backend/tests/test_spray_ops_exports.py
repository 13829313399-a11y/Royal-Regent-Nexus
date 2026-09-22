from io import BytesIO
from decimal import Decimal

import openpyxl
from sqlalchemy.orm import Session

from test_spray_ops import client, payload, post, read, user_with, setup_order, start, publish, manual_scenario, report
from app.models import spray_ops as m


def test_export_all_filtered_rows_freezes_and_preserves_identifiers(client):
    # More than the API's default page size; no dependence on pagination.
    with Session(client.spray_engine) as db:
        for i in range(65):
            batch = m.SprayOpsBatch(factory_id="huaxing", document_no=f"B{i:04d}", item_no="00017", part="左", quantity=1, accepted=1, held=0, rejected=0, unit="件", business_date="2026-09-02")
            db.add(batch)
            db.flush()
            db.add(m.SprayOpsStock(factory_id="huaxing", batch_id=batch.id, state="white", quantity=1, reserved=0, unit="件", ready_at="2026-09-02T00:00:00+00:00"))
        db.commit()
    body = payload(kind="stock", base_revision=0)
    result = client.post("/api/spray-operations/exports", json=body)
    assert result.status_code == 200, result.text
    artifact = result.json()["data"]
    link = f"/api/spray-operations/exports/{artifact['id']}/download"
    first = client.get(link, params={"factory_id": "huaxing"})
    book = openpyxl.load_workbook(BytesIO(first.content))
    sheet = book.active
    records = list(sheet.iter_rows(min_row=4, max_row=68, values_only=True))
    assert len(records) == 65 and all(row[2] == "00017" for row in records)
    assert sheet['C4'].number_format == '@'
    assert sheet.page_setup.orientation == "landscape" and sheet.print_title_rows == "$1:$3"
    assert client.get(link, params={"factory_id": "huaxing"}).content == first.content
    assert client.post("/api/spray-operations/exports", json=body).json()["data"] == artifact
    client.spray_auth["user"] = user_with("read")
    assert client.get(link, params={"factory_id": "huaxing"}).status_code == 403


def test_daily_print_seven_people_never_multiplies_team_production(client):
    setup = setup_order(client, quantity='560')
    task = start(client, publish(client, manual_scenario(client, setup, qty='560')))
    people = [post(client, 'employees', code=f'PRINT-{i}', name=f'合成工人{i}') for i in range(7)]
    daily = report(client, task, qty='560', good='560', hold='0', scrap='0')
    post(client, f"reports/{daily['id']}/labor", expected_version=daily['version'], row_id=daily['rows'][0]['id'], business_date='2026-09-02', reason='合成签认', labor=[dict(employee_id=p['id'], hours='8') for p in people])
    revision = client.get('/api/spray-operations/access', params=dict(factory_id='huaxing')).json()['meta']['factory_revision']
    artifact = post(client, 'exports', kind='daily', base_revision=revision, from_date='2026-09-02', to_date='2026-09-02')
    link = f"/api/spray-operations/exports/{artifact['id']}/download"
    book = openpyxl.load_workbook(BytesIO(client.get(link, params=dict(factory_id='huaxing')).content))
    printed = list(book.active.iter_rows(min_row=4, max_row=4, values_only=True))
    assert sum(Decimal(row[6] or 0) for row in printed) == 560
    assert printed[0][12] == 7
    assert book.active['F4'].number_format == '#,##0'
    assert not book.active['F4'].alignment.wrap_text
    workers = list(book['员工工时与计薪'].iter_rows(min_row=4, max_row=10, values_only=True))
    assert len({row[5] for row in workers}) == 7
    assert sum(Decimal(row[6]) for row in workers) == 56
    assert sum(book.active.column_dimensions[col].width for col in 'ABCDEFGHIJKLMN') < 150
    # The frozen export was produced with payroll columns, so losing that right
    # also revokes this artifact even when general export remains granted.
    client.spray_auth['user'] = user_with('read', 'export')
    assert client.get(link, params=dict(factory_id='huaxing')).status_code == 403
