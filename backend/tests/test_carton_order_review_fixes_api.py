from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from threading import Barrier
from uuid import uuid4
import json
import os
import sqlite3
from io import BytesIO

from openpyxl import load_workbook
import pytest

from test_molding_sample_api import make_client, login_as
from test_carton_procurement_api import _ensure_dickie_customer, _order_payload, _submit_order, _freeze_carton_time

BASE = "/api/carton-procurement"


def create(client, **changes):
    payload = {**_order_payload(), **changes}
    response = client.post(BASE + "/orders", json=payload)
    assert response.status_code == 201, response.text
    return response.json()


def context(client, row):
    return client.get(f"{BASE}/orders/{row['order_no']}/purchase-order-context", params={"factory_id": "huaxing"}).json()


def test_new_order_request_replays_original_result_survives_changes_and_allows_intentional_blank_po_orders(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "warehouse_keeper")
        _ensure_dickie_customer(client)
        _freeze_carton_time(monkeypatch)
        payload = {**_order_payload(), "request_id": uuid4().hex}
        barrier = Barrier(2)

        def send():
            barrier.wait(timeout=10)
            return client.post(BASE + "/orders", json=payload)

        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = [future.result(timeout=30) for future in [pool.submit(send), pool.submit(send)]]
        assert [response.status_code for response in responses] == [201, 201]
        first = responses[0].json()
        assert responses[1].json() == first
        # The original reply is durable even after a subsequent workflow transition.
        _submit_order(client, first)
        assert client.post(BASE + "/orders", json=payload).json() == first
        numeric_equivalent = {**payload, "product_order_quantity": "3600.000000"}
        assert client.post(BASE + "/orders", json=numeric_equivalent).json() == first
        changed = {**payload, "note": "超时后修改了内容"}
        assert client.post(BASE + "/orders", json=changed).status_code == 409
        second = create(client, request_id=uuid4().hex)
        assert second["id"] != first["id"]
        assert second["customer_po"] == first["customer_po"] == ""
        orders = client.get(BASE + "/orders", params={"factory_id": "huaxing"}).json()
        assert orders["total"] == 2
        audit = client.get(BASE + "/audit-events", params={"factory_id": "huaxing", "event_type": "ORDER_CREATED"})
        assert audit.status_code == 200, audit.text
        assert all("request" not in item["detail"] for item in audit.json()["items"])
        login_as(client, "admin")
        another_actor = client.post(BASE + "/orders", json=payload)
        assert another_actor.status_code == 201, another_actor.text
        assert another_actor.json()["id"] not in {first["id"], second["id"]}


def test_failed_order_request_does_not_reserve_request_id(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "warehouse_keeper")
        _ensure_dickie_customer(client)
        request_id = uuid4().hex
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonImportBatch, CartonAuditEvent
        from app.services.carton_schedule_tracking import identity
        source = {"schedule_customer_code": "DICKIE", "customer_name": "Dickie", "contract_no": "SC700145365",
                  "item_no": "203302044", "quantity": 3600, "order_type": "正单", "schedule_section": "PENDING",
                  "source_sheet": "ITEM", "source_row": 12}
        source["schedule_identity"] = identity(source)
        batch_id = "CIB-" + uuid4().hex
        with SessionLocal() as db:
            db.add(CartonImportBatch(id=batch_id, factory_id="huaxing", import_type="WEEKLY_SCHEDULE",
                original_filename="schedule.xlsx", source_sha256=uuid4().hex, status="CONFIRMED",
                parse_summary_json=json.dumps({"rows": [source]}), imported_by="user-warehouse-keeper",
                imported_by_name="Warehouse", created_at="2026-08-05T10:00:00+08:00"))
            db.add(CartonAuditEvent(id="CAE-" + uuid4().hex, factory_id="huaxing", event_type="IMPORT_BATCH_CREATED",
                entity_type="carton_import_batch", entity_id=batch_id, detail_json="{}", actor_user_id="user-warehouse-keeper",
                actor_name="Warehouse", created_at="2026-08-05T10:00:00+08:00"))
            db.commit()
        payload = {**_order_payload(), "request_id": request_id,
                   "schedule_source": {"batch_id": batch_id, "source_sheet": "ITEM", "source_row": 12}}
        payload["lines"] = [{**line, "paper_quality": ""} for line in payload["lines"]]
        assert client.post(BASE + "/orders", json=payload).status_code == 422
        with SessionLocal() as db:
            from sqlalchemy import select
            assert not db.scalars(select(CartonAuditEvent).where(CartonAuditEvent.event_type == "SCHEDULE_ORDER_LINKED")).all()
        payload["lines"] = _order_payload()["lines"]
        first = client.post(BASE + "/orders", json=payload)
        assert first.status_code == 201, first.text
        row = first.json()
        assert row["revision"] == 1
        assert client.post(BASE + "/orders", json=payload).json() == row
        assert client.get(BASE + "/orders", params={"factory_id": "huaxing"}).json()["total"] == 1
        with SessionLocal() as db:
            assert len(db.scalars(select(CartonAuditEvent).where(CartonAuditEvent.event_type == "SCHEDULE_ORDER_LINKED")).all()) == 1
            assert len(db.scalars(select(CartonAuditEvent).where(CartonAuditEvent.event_type == "ORDER_CREATED")).all()) == 1


def test_date_authorization_reads_state_only_after_factory_lock(monkeypatch):
    monkeypatch.setenv("AUTHZ_MODE", "enforce")
    with make_client(monkeypatch) as client:
        login_as(client, "warehouse_keeper")
        _ensure_dickie_customer(client)
        row = create(client)
        from app.services import carton_procurement as core
        from app.services.auth import AuthContext, AuthGrantContext, get_current_user
        grant = AuthGrantContext("test-role", "Test role", "huaxing", "pmc-warehouse", frozenset({"carton_procurement:order_write"}))
        actor = AuthContext("user-warehouse-keeper", "scoped", "Scoped", (), (), grant.permissions,
                            ("huaxing",), ("pmc-warehouse",), grants=(grant,))
        client.app.dependency_overrides[get_current_user] = lambda: actor
        original_lock = core._lock_receipt_factory
        transition_observed = False

        def observe_concurrent_confirmation(db, factory):
            nonlocal transition_observed
            original_lock(db, factory)
            if not transition_observed:
                transition_observed = True
                # A confirmation won the factory lock before this date request.
                current = core.get_order_by_no(db, factory, row["order_no"])
                current.status = "PENDING_SUPPLIER"
                current.revision = row["revision"] + 1
                db.flush()

        monkeypatch.setattr(core, "_lock_receipt_factory", observe_concurrent_confirmation)
        response = client.post(f"{BASE}/orders/{row['order_no']}/dates", json={
            "factory_id": "huaxing", "expected_revision": row["revision"] + 1,
            "due_date": "2026-08-20", "reason": "并发确认后调整交期"})
        assert response.status_code == 403, response.text
        assert transition_observed is True
        client.app.dependency_overrides.clear()
        current = client.get(BASE + "/orders", params={"factory_id": "huaxing"}).json()["items"][0]
        assert current["due_date"] == row["due_date"]


@pytest.mark.parametrize("authorization_mode", ["legacy", "shadow", "enforce"])
def test_date_changes_enforce_scoped_canonical_permissions_and_denies_in_every_mode(monkeypatch, authorization_mode):
    monkeypatch.setenv("AUTHZ_MODE", authorization_mode)
    with make_client(monkeypatch) as client:
        login_as(client, "warehouse_keeper")
        _ensure_dickie_customer(client)
        _freeze_carton_time(monkeypatch)
        confirmed = create(client)
        locked = _submit_order(client, create(client))
        from app.services.auth import AuthContext, AuthGrantContext, AuthOverrideContext, get_current_user

        def grant(permission, factory="huaxing", department="pmc-warehouse"):
            return AuthGrantContext("test-role", "Test role", factory, department,
                                    frozenset({f"carton_procurement:{permission}"}))

        def use_actor(grants, overrides=()):
            actor = AuthContext("user-warehouse-keeper", "scoped", "Scoped", (), (),
                frozenset(p for g in grants for p in g.permissions), ("huaxing",),
                ("pmc-warehouse",), grants=tuple(grants), overrides=tuple(overrides))
            client.app.dependency_overrides[get_current_user] = lambda: actor

        denied_cases = [
            (confirmed, [grant("order_write")], "order_write"),
            (confirmed, [grant("order_write", "huadeng")], None),
            (confirmed, [grant("order_write", department="qc")], None),
            (locked, [grant("order_write")], None),
            (locked, [grant("order_write"), grant("inventory_write", "huadeng")], None),
            (locked, [grant("order_write"), grant("inventory_write", department="qc")], None),
            (locked, [grant("order_write"), grant("inventory_write")], "inventory_write"),
            (locked, [grant("order_write"), grant("inventory_write")], "order_write"),
            (locked, [grant("order_write"), grant("order_adjust")], "order_adjust"),
        ]
        for row, grants, denied_permission in denied_cases:
            overrides = [AuthOverrideContext("deny", f"carton_procurement:{denied_permission}",
                         "deny", "huaxing", "pmc-warehouse")] if denied_permission else []
            use_actor(grants, overrides)
            response = client.post(f"{BASE}/orders/{row['order_no']}/dates", json={
                "factory_id": "huaxing", "expected_revision": row["revision"],
                "due_date": "2026-08-22", "reason": "核验交期调整权限边界"})
            assert response.status_code == 403, (authorization_mode, denied_permission, response.text)
        client.app.dependency_overrides.clear()
        current = {row["id"]: row for row in client.get(BASE + "/orders", params={"factory_id": "huaxing"}).json()["items"]}
        for row in (confirmed, locked):
            assert (current[row["id"]]["due_date"], current[row["id"]]["revision"]) == (row["due_date"], row["revision"])
        audit = client.get(BASE + "/audit-events", params={"factory_id": "huaxing", "event_type": "ORDER_DATES_CHANGED"})
        assert audit.status_code == 200 and audit.json()["total"] == 0

        use_actor([grant("order_write")])
        response = client.post(f"{BASE}/orders/{confirmed['order_no']}/dates", json={
            "factory_id": "huaxing", "expected_revision": confirmed["revision"],
            "due_date": "2026-08-22", "reason": "待下单订单正常调整"})
        assert response.status_code == 200, response.text
        for grants, overrides, target in (
            ([grant("order_write"), grant("inventory_write")], [], "2026-08-23"),
            ([grant("order_write"), grant("order_adjust")],
             [AuthOverrideContext("deny", "carton_procurement:inventory_write", "deny", "huaxing", "pmc-warehouse")], "2026-08-24"),
        ):
            use_actor(grants, overrides)
            response = client.post(f"{BASE}/orders/{locked['order_no']}/dates", json={
                "factory_id": "huaxing", "expected_revision": locked["revision"],
                "due_date": target, "reason": "已锁定订单授权调整"})
            assert response.status_code == 200, response.text
            locked = response.json()
        client.app.dependency_overrides.clear()


def test_date_only_change_preserves_quantity_lead_and_explicit_dates_and_rejects_stale_or_terminal(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "warehouse_keeper")
        _ensure_dickie_customer(client)
        _freeze_carton_time(monkeypatch)
        row = _submit_order(client, create(client, customer_due_date="2026-08-15"))
        with sqlite3.connect(os.environ["DATABASE_URL"].removeprefix("sqlite:///")) as db:
            db.execute("UPDATE carton_orders SET safety_lead_days=5 WHERE id=?", (row["id"],))
        request = {"factory_id": "huaxing", "expected_revision": row["revision"],
                   "customer_due_date": "2026-08-17", "due_date": "2026-08-25", "reason": "客人确认新的交期"}
        response = client.post(f"{BASE}/orders/{row['order_no']}/dates", json=request)
        assert response.status_code == 200, response.text
        updated = response.json()
        assert updated["due_date"] == "2026-08-12"  # Same plan, changed customer deadline still creates C.
        assert updated["safety_lead_days"] == 5
        assert updated["lines"] == row["lines"]
        assert updated["product_order_quantity"] == row["product_order_quantity"]
        assert context(client, updated)["pending_type"] == "ADJUSTMENT"
        assert client.post(f"{BASE}/orders/{row['order_no']}/dates", json=request).status_code == 409
        issue = client.post(f"{BASE}/orders/{row['order_no']}/purchase-order-issues.xlsx", json={
            "factory_id": "huaxing", "expected_revision": updated["revision"]})
        assert issue.status_code == 200 and issue.headers["x-purchase-order-document-no"].endswith("-C01")
        assert context(client, updated)["pending_type"] == "NONE"
        assert client.post(f"{BASE}/orders/{row['order_no']}/dates", json={
            **request, "expected_revision": updated["revision"], "customer_due_date": "2026-08-04"}).status_code == 422
        barrier = Barrier(2)

        def change_date(day):
            barrier.wait(timeout=10)
            return client.post(f"{BASE}/orders/{row['order_no']}/dates", json={
                **request, "expected_revision": updated["revision"], "customer_due_date": day})

        with ThreadPoolExecutor(max_workers=2) as pool:
            changes = [future.result(timeout=30) for future in [pool.submit(change_date, "2026-08-19"), pool.submit(change_date, "2026-08-20")]]
        assert sorted(response.status_code for response in changes) == [200, 409]
        explicit = create(client, quantity_basis="EXPLICIT", product_order_quantity=None, lines=[{
            **_order_payload()["lines"][0], "usage_quantity": None, "required_quantity": "20"}])
        explicit = _submit_order(client, explicit)
        response = client.post(f"{BASE}/orders/{explicit['order_no']}/dates", json={
            "factory_id": "huaxing", "expected_revision": explicit["revision"],
            "due_date": "2026-08-25", "customer_due_date": "2026-08-28", "reason": "历史明确需求调整计划"})
        assert response.status_code == 200, response.text
        updated_explicit = response.json()
        assert updated_explicit["due_date"] == "2026-08-25"
        assert updated_explicit["lines"] == explicit["lines"]
        cancelled = client.post(f"{BASE}/orders/{explicit['order_no']}/reduce", json={
            "factory_id": "huaxing", "expected_revision": updated_explicit["revision"],
            "line_quantities": [{"order_line_id": explicit["lines"][0]["id"], "required_quantity": "0"}]}).json()
        assert cancelled["status"] == "CANCELLED"
        assert client.post(f"{BASE}/orders/{explicit['order_no']}/dates", json={
            "factory_id": "huaxing", "expected_revision": cancelled["revision"],
            "due_date": "2026-08-26", "reason": "结束订单禁止修改"}).status_code == 409


def test_full_reduction_automatically_issues_zero_target_cancellation_without_rewriting_original_quantities(monkeypatch):
    with make_client(monkeypatch) as client:
        login_as(client, "warehouse_keeper")
        _ensure_dickie_customer(client)
        _freeze_carton_time(monkeypatch)
        row = _submit_order(client, create(client, net_weight_kg="1.5", gross_weight_kg="2"))
        # Deactivated master data must not block an already-issued order's cancellation.
        with sqlite3.connect(os.environ["DATABASE_URL"].removeprefix("sqlite:///")) as db:
            db.execute("UPDATE carton_customers SET status='INACTIVE' WHERE factory_id='huaxing' AND customer_code='DICKIE'")
        response = client.post(f"{BASE}/orders/{row['order_no']}/reduce", json={
            "factory_id": "huaxing", "expected_revision": row["revision"], "reduction_quantity": "3600"})
        assert response.status_code == 200, response.text
        cancelled = response.json()
        assert cancelled["status"] == "CANCELLED"
        assert cancelled["product_order_quantity"] == row["product_order_quantity"]
        assert cancelled["lines"] == row["lines"]
        result = context(client, cancelled)
        assert result["pending_type"] == "NONE"
        cancellation = result["issues"][0]
        assert cancellation["is_cancellation"] is True
        assert cancellation["document_type"] == "REDUCE"
        assert cancellation["document_no"].endswith("-R01")
        assert Decimal(cancellation["after_product_quantity"]) == 0
        with sqlite3.connect(os.environ["DATABASE_URL"].removeprefix("sqlite:///")) as db:
            snapshot = json.loads(db.execute("SELECT snapshot_json FROM carton_purchase_order_issues WHERE id=?", (cancellation["id"],)).fetchone()[0])
            assert snapshot["cancellation"] is True
            assert all(Decimal(line["after_required_quantity"]) == 0 for line in snapshot["lines"])
            # A legacy cancellation without its R document exposes a repairable pending notice.
            db.execute("DELETE FROM carton_purchase_order_issues WHERE id=?", (cancellation["id"],))
        assert context(client, cancelled)["can_generate"] is True
        repaired = client.post(f"{BASE}/orders/{row['order_no']}/purchase-order-issues.xlsx", json={
            "factory_id": "huaxing", "expected_revision": cancelled["revision"]})
        assert repaired.status_code == 200, repaired.text
        workbook = load_workbook(BytesIO(repaired.content))
        sheet = workbook.active
        assert sheet.title == "取消采购通知"
        assert "取消采购通知" in sheet["A1"].value
        assert sheet["J10"].value == "停止后续生产和送货"
        assert "停止后续生产和送货" in sheet.cell(row=11 + len(row["lines"]), column=1).value
        workbook.close()
        from app.db import SessionLocal
        from app.models.carton_procurement import CartonPurchaseOrderIssue
        from app.services.carton_procurement_export import build_purchase_order_issue_batch_workbook
        from app.services.carton_procurement import business_now
        with SessionLocal() as db:
            from sqlalchemy import select
            issue = db.scalar(select(CartonPurchaseOrderIssue).where(CartonPurchaseOrderIssue.id == repaired.headers["x-purchase-order-issue-id"]))
            workbook = load_workbook(BytesIO(build_purchase_order_issue_batch_workbook([issue], generated_at=business_now())))
            assert workbook.active["C6"].value == "取消采购通知"
            assert "停止后续生产和送货" in workbook.active.cell(row=7 + len(row["lines"]), column=1).value
            workbook.close()
        with sqlite3.connect(os.environ["DATABASE_URL"].removeprefix("sqlite:///")) as db:
            db.execute("UPDATE carton_customers SET status='ACTIVE' WHERE factory_id='huaxing' AND customer_code='DICKIE'")
        blank = create(client)
        response = client.post(f"{BASE}/orders/{blank['order_no']}/cancel", json={
            "factory_id": "huaxing", "expected_revision": blank["revision"], "reason": "未下单直接取消"})
        assert response.status_code == 200
        assert context(client, response.json())["can_generate"] is False


def test_order_return_canonical_permissions_prevent_inventory_deduction_and_authorized_return_issues_notice(monkeypatch):
    # Explicit deny must hold in legacy rollout too, without an inventory side effect.
    monkeypatch.setenv("AUTHZ_MODE", "legacy")
    with make_client(monkeypatch) as client:
        login_as(client, "warehouse_keeper")
        _ensure_dickie_customer(client)
        _freeze_carton_time(monkeypatch)
        row = _submit_order(client, create(client))
        receipt = client.post(BASE + "/receipts", json={
            "factory_id": "huaxing", "delivery_note_no": uuid4().hex, "delivery_date": "2026-08-05",
            "lines": [{"order_line_id": row["lines"][0]["id"], "delivered_quantity": "5",
                       "received_quantity": "5", "unit_price": "2", "location": "A-01"}]}).json()
        response = client.post(f"{BASE}/receipts/{receipt['id']}/confirm", json={
            "factory_id": "huaxing", "expected_revision": receipt["revision"]})
        assert response.status_code == 200, response.text
        row = client.get(BASE + "/orders", params={"factory_id": "huaxing"}).json()["items"][0]
        from app.services.auth import AuthContext, AuthGrantContext, AuthOverrideContext, get_current_user

        def grant(permission, factory="huaxing", department="pmc-warehouse"):
            return AuthGrantContext("test-role", "Test role", factory, department, frozenset({f"carton_procurement:{permission}"}))

        for case in ("order_only", "inventory_only", "foreign_inventory", "other_department", "deny_inventory", "deny_order"):
            grants = [grant("inventory_write")] if case == "inventory_only" else [grant("order_write")]
            if case not in {"order_only", "inventory_only"}:
                grants.append(grant("inventory_write", "huadeng" if case == "foreign_inventory" else "huaxing",
                                    "qc" if case == "other_department" else "pmc-warehouse"))
            overrides = (AuthOverrideContext("deny", f"carton_procurement:{case.removeprefix('deny_')}_write", "deny", "huaxing", "pmc-warehouse"),) if case.startswith("deny_") else ()
            actor = AuthContext("user-warehouse", "scoped", "Scoped", (), (), frozenset(p for g in grants for p in g.permissions),
                                ("huaxing",), ("pmc-warehouse",), grants=tuple(grants), overrides=overrides)
            client.app.dependency_overrides[get_current_user] = lambda: actor
            response = client.post(f"{BASE}/orders/{row['order_no']}/return", json={
                "factory_id": "huaxing", "expected_revision": row["revision"], "reason": "客户取消整单退回"})
            assert response.status_code == 403, (case, response.text)
        client.app.dependency_overrides.clear()
        ledger = client.get(BASE + "/inventory/movements", params={"factory_id": "huaxing"}).json()["items"]
        assert sum(Decimal(item["quantity"]) for item in ledger) == 5
        response = client.post(f"{BASE}/orders/{row['order_no']}/return", json={
            "factory_id": "huaxing", "expected_revision": row["revision"], "reason": "客户取消整单退回"})
        assert response.status_code == 200, response.text
        assert response.json()["status"] == "CANCELLED"
        assert context(client, response.json())["issues"][0]["is_cancellation"] is True
        ledger = client.get(BASE + "/inventory/movements", params={"factory_id": "huaxing"}).json()["items"]
        assert sum(Decimal(item["quantity"]) for item in ledger) == 0
