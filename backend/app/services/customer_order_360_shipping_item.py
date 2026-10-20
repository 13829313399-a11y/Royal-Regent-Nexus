"""360's shipping item is a display identifier, independent of the full PO item."""
from __future__ import annotations

from copy import copy
import re

from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter


HEADER = "出货货号"
COMPLETED = "已完成（按验货日期）"
ISSUE_CODE = "huakang_a_360_shipping_item"
_BOUNDARIES = {"取消单", "已走货订单", "已走货", "已完成", COMPLETED}


class ShippingItemError(ValueError):
    pass


def _text(value):
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def derive_shipping_item(value) -> str:
    """Apply only the three confirmed numeric shapes; never coerce short strings."""
    text = _text(value)
    if re.fullmatch(r"[0-9]{11}", text):
        return text[:7]
    if re.fullmatch(r"[0-9]{10}\.[0-9]+", text):
        return text.split(".", 1)[0]
    return text if re.fullmatch(r"[0-9]{7}|[0-9]{10}", text) else ""


def _get(record, field, default=""):
    return record.get(field, default) if isinstance(record, dict) else getattr(record, field, default)


def _key(value):
    # Do not strip a full product's suffix, decimal point or leading zeroes.
    return _text(value).casefold()


def _history_value(item):
    return _text((_get(item, "extras", {}) or {}).get("shipping_item"))


def _history_source(items):
    return "、".join(dict.fromkeys(
        f"{_get(item, 'source_sheet', 'ITEM表')}第{_get(item, 'row', '?')}行"
        for item in items
    ))


def _resolve(product, history, reference=None):
    matches = [item for item in history if _key(_get(item, "product_no")) == _key(product)
               and _history_value(item)] if _text(product) else []
    match_label = "同完整货号"
    if reference:
        same_order = [item for item in matches if _key(_get(item, "reference_no")) == _key(reference)]
        if same_order:
            matches = same_order
            match_label = "同Release及完整货号"
    values = {_history_value(item) for item in matches}
    if len(values) > 1:
        return "", "conflict", f"{_history_source(matches)}{match_label}的出货货号不一致，留空核对"
    if values:
        return next(iter(values)), "history", f"{_history_source(matches)} · {match_label}的唯一历史出货货号"
    value = derive_shipping_item(product)
    if value:
        return value, "rule", "旧排期格式规则：11位数字取前7位；10位数字的小数尾去除；7/10位数字保留"
    return "", "unknown", "完整产品编号不属于已确认的出货货号格式，留空核对"


def prepare_shipping_item(row, history):
    """Recompute after a product correction, retaining a single optional edit issue."""
    value, _, source = _resolve(row.get("product_no"), history, row.get("reference_no"))
    row["shipping_item"] = value
    row.setdefault("lineage", {})["shipping_item"] = source
    issues = row.setdefault("issues", [])
    issues[:] = [issue for issue in issues if issue.get("code") != ISSUE_CODE]
    issues.append({
        "severity": "warning", "code": ISSUE_CODE, "field": "shipping_item",
        "message": f"出货货号：{source}。可按业务需要补填或修改；不影响完整产品编号。",
        "can_skip": True, "skip_key": f"{row.get('id', '')}|{ISSUE_CODE}|shipping_item",
        "skip_label": "", "can_edit": True, "edit_field": "shipping_item",
        "edit_label": HEADER, "edit_input_type": "text",
    })
    status = "blocked" if any(issue.get("severity") == "blocked" for issue in issues) else "warning"
    row["status"] = status
    row["status_label"] = {"blocked": "已阻断", "warning": "需复核"}[status]


def _enabled(factory_id, customer_code):
    return factory_id == "huakang-a" and customer_code == "360"


def _header(value):
    return re.sub(r"\s+", "", _text(value))


def _occupied(sheet, column):
    if any(c == column and (cell.value is not None or cell.data_type == "f" or cell.comment is not None)
           for (_, c), cell in sheet._cells.items()):
        return True
    if any(area.min_col <= column <= area.max_col for area in sheet.merged_cells.ranges):
        return True
    return any(area.min_col <= column <= area.max_col
               for rule in sheet.data_validations.dataValidation for area in rule.sqref.ranges)


def ensure_shipping_item_column(workbook, factory_id="huakang-a", customer_code="360") -> int | None:
    """Use a free AM, otherwise append without shifting any historical column."""
    if not _enabled(factory_id, customer_code):
        return None
    if "ITEM表" not in workbook.sheetnames:
        raise ShippingItemError("360 出货货号：缺少 ITEM表，未创建列。")
    sheet = workbook["ITEM表"]
    found = [column for (row, column), cell in sheet._cells.items()
             if row == 3 and column > 28 and cell.data_type != "f" and _header(cell.value) == HEADER]
    if len(found) > 1:
        raise ShippingItemError("360 出货货号：ITEM表存在多个“出货货号”表头，请先保留唯一列。")
    if found:
        return found[0]
    column = 39
    if _occupied(sheet, column):
        # A validation alone does not increase openpyxl's max_column.
        column = max(sheet.max_column + 1, 40)
        while column <= 16_384 and _occupied(sheet, column):
            column += 1
    if column > 16_384:
        raise ShippingItemError("360 出货货号：没有可安全追加的空列，请先核对 ITEM表。")
    header = sheet.cell(3, column)
    header._style = copy(sheet.cell(3, 8)._style)
    header.value = HEADER
    header.data_type = "s"
    header.number_format = "@"
    alignment = copy(header.alignment)
    alignment.wrap_text = True
    header.alignment = alignment
    sheet.column_dimensions[get_column_letter(column)].width = 14
    return column


def _literal(sheet, row, column):
    cell = sheet._cells.get((row, column))
    return _text(cell.value) if cell is not None and cell.data_type != "f" else ""


def _order_identity(sheet, row):
    reference = _literal(sheet, row, 5) or _literal(sheet, row, 6) or _literal(sheet, row, 4)
    product = _literal(sheet, row, 8)
    if _header(product).casefold() in {"产品编号", "产品货号", "货号", "itemno", "item#"}:
        return "", ""
    return reference, product


def _rows_and_boundaries(sheet):
    cells = list(sheet._cells.values())
    simulated = {cell.row for cell in cells if isinstance(cell.value, str) and "模拟数据" in cell.value}
    markers = sorted((cell.row, _text(cell.value)) for cell in cells
                     if cell.column == 1 and cell.row >= 4 and cell.data_type != "f"
                     and _text(cell.value) in _BOUNDARIES and not all(_order_identity(sheet, cell.row)))
    completed = [row for row, label in markers if label == COMPLETED]
    if len(completed) > 1:
        raise ShippingItemError("360 出货货号：ITEM表存在多个按验货日期已完成分区，请先核对标题。")
    stop = min((row for row, _ in markers), default=sheet.max_row + 1)
    start = completed[0] if completed else None
    end = min((row for row, _ in markers if start is not None and row > start), default=sheet.max_row + 1)
    rows = sorted({cell.row for cell in cells if cell.row >= 4 and cell.value is not None
                   and cell.row not in simulated and all(_order_identity(sheet, cell.row))})
    targets = [row for row in rows if row < stop or (start is not None and start < row < end)]
    return rows, targets


def backfill_shipping_items(workbook, factory_id="huakang-a", customer_code="360") -> dict:
    """Fill blank pending/explicitly date-completed rows; never overwrite history."""
    report = {"applied": False, "column": None, "filled": 0, "preserved": 0,
              "unresolved": [], "warnings": []}
    if not _enabled(factory_id, customer_code):
        return report
    if "ITEM表" not in workbook.sheetnames:
        raise ShippingItemError("360 出货货号：缺少 ITEM表，未补填。")
    sheet = workbook["ITEM表"]
    # Validate the section layout before adding a column or writing any values.
    rows, targets = _rows_and_boundaries(sheet)
    column = ensure_shipping_item_column(workbook, factory_id, customer_code)
    report.update(applied=True, column=get_column_letter(column))
    history = []
    for row in rows:
        value = _literal(sheet, row, column)
        if value:
            reference, product = _order_identity(sheet, row)
            history.append({"row": row, "source_sheet": sheet.title, "reference_no": reference,
                            "product_no": product, "extras": {"shipping_item": value}})
    plans = []
    for row in targets:
        current = sheet._cells.get((row, column))
        if current is not None and (current.value not in (None, "") or current.data_type == "f"):
            report["preserved"] += 1
            continue
        reference, product = _order_identity(sheet, row)
        coordinate = f"{get_column_letter(column)}{row}"
        if any(area.min_row <= row <= area.max_row and area.min_col <= column <= area.max_col
               for area in sheet.merged_cells.ranges):
            raise ShippingItemError(f"360 出货货号：ITEM表!{coordinate} 位于合并区域，未覆盖。")
        value, reason, source = _resolve(product, history)
        if not value:
            detail = {"sheet": sheet.title, "row": row, "cell": coordinate,
                      "reference_no": reference, "product_no": product, "reason": reason}
            report["unresolved"].append(detail)
            report["warnings"].append(f"ITEM表!{coordinate}（{reference} / {product}）：{source}。")
            continue
        plans.append((row, value))
    for row, value in plans:
        cell = sheet.cell(row, column)
        if isinstance(cell, MergedCell):  # Defensive; all targets were checked above.
            raise ShippingItemError(f"360 出货货号：{cell.coordinate} 是合并单元格，未写入。")
        if not cell.has_style:
            cell._style = copy(sheet.cell(row, 8)._style)
        alignment = copy(cell.alignment)
        alignment.wrap_text = True
        cell.alignment = alignment
        cell.value = value
        cell.data_type = "s"
        cell.number_format = "@"
    report["filled"] = len(plans)
    return report
