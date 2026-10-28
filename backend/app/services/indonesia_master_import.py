"""Preview-first material master intake; shipment facts never become BOM usage."""
from collections import defaultdict
from copy import deepcopy
from io import BytesIO
from pathlib import PurePath
import hashlib
import re
from xml.etree import ElementTree
from zipfile import ZipFile

import openpyxl
from openpyxl.utils.cell import range_boundaries
from fastapi import HTTPException
from app.schemas import indonesia_materials as s
from app.services.indonesia_material_calculations import require
from app.services.indonesia_materials import Ledger, authorize, digest, view
from app.services.indonesia_warehouse_import import normalized, text, value_key
from app.services.indonesia_master_workbook import template_units, formula_cache, sheet_photos, NUMERIC_LABELS
from app.services.identity_policy import lock_mutation
from app.services.transaction_lock import lock_transaction

TYPE = 'material_master'
IDENTITY_RULE = 'name-spec-v2'
FORMAT_RULE = 'unified-template-v1'
MAX_ROWS = 2000
FIELDS = {
    'item_no': ['货号', '成品货号', '产品货号', 'Item No.', 'Item No'],
    'code': ['物料编码', '物料编号', '料号'],
    'name': ['正式报关名称', '正式报关中文名称', '物料名称', '产品中文名称', '品名'],
    'english_name': ['英文名称', '产品英文名称', '物料英文名称'],
    'specification': ['规格', '规格型号'], 'category': ['类别', '物料类别'],
    'customs_unit': ['报关单位'], 'unit': ['数量单位', '计数单位'],
    'price': ['采购单价', '单价'], 'currency': ['币种'], 'supplier': ['供应商', '供应商名称'],
    'usage': ['产品用途', '用途'], 'department': ['使用部门'],
    'export_company': ['出口公司', '报关出口公司'],
    'net_weight_kg': ['单个净重', '单个净重kg', '单个净重KGM'],
    'procurement_days': ['采购备料自然日', '采购周期', '采购周期天'],
    'piece_usage': ['单件用量', '每件成品用量', '单件物料用量'],
    'image_data': ['图片', '物料图片', '产品图片', 'Picture'],
}
MATERIAL_FIELDS = set(FIELDS) - {'item_no', 'piece_usage'}
ALIASES = {normalized(alias): key for key, names in FIELDS.items() for alias in [key, *names]}


def header_field(value):
    # Chinese/English bilingual headers commonly occupy a single cell.
    first = str(value).strip().split('\n')[0].strip()
    return ALIASES.get(normalized(value)) or ALIASES.get(normalized(first))


def parse_material_book(content, filename, mappings=None):
    """Runs in the existing bounded document worker, without access to the ledger."""
    require(filename.lower().endswith('.xlsx'), '请上传统一资料 .xlsx 文件', 422)
    require(0 < len(content) <= 25 * 1024 * 1024, '文件为空或超过 25 MB', 413)
    book = cached_book = None
    try:
        with ZipFile(BytesIO(content)) as archive:
            require(sum(item.file_size for item in archive.infolist()) <= 150 * 1024 * 1024, '文件解压后过大，请分批导入', 422)
            book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
            cached_book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
            require(len(book.worksheets) <= 30, '工作表超过 30 张，请分批导入', 422)
            overrides = {item.sheet: item for item in mappings or []}
            require(len(overrides) == len(mappings or []), '工作表映射不能重复', 422)
            require(set(overrides) <= set(book.sheetnames), '映射中的工作表不存在', 422)
            result, sheets, issues = [], [], []
            image_budget = {'bytes': 0}
            for sheet in book:
                # Irrelevant invoice sheets can have inflated formatting dimensions.
                heads = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 1, 30), max_col=120))
                index, columns = None, {}
                override = overrides.get(sheet.title)
                if override:
                    index, columns = override.header_row - 1, dict(override.columns)
                    require(index < len(heads), '表头行超出工作表范围', 422)
                else:
                    for position, cells in enumerate(heads):
                        found = defaultdict(list)
                        for col, cell in enumerate(cells, 1):
                            field = header_field(text(cell))
                            if field:
                                found[field].append(col)
                        if 'item_no' in found and ('name' in found or 'code' in found):
                            index = position
                            columns = {field: indexes[0] for field, indexes in found.items() if len(indexes) == 1}
                            columns.update(template_units([text(cell) for cell in cells], normalized))
                            break
                if index is None:
                    continue
                require(set(columns) <= set(FIELDS), '包含不支持的字段映射', 422)
                require(len(set(columns.values())) == len(columns), '同一来源列不能映射到多个字段', 422)
                headers = [text(cell)[:300] for cell in heads[index]]
                while headers and not headers[-1]:
                    headers.pop()
                require(all(1 <= col <= len(headers) for col in columns.values()), '来源列超出表头范围', 422)
                sheets.append(dict(sheet=sheet.title, header_row=index + 1, headers=headers, columns=columns,
                                   enabled=override.enabled if override else True))
                if override and not override.enabled:
                    continue
                if not columns.get('item_no') or not (columns.get('name') or columns.get('code')):
                    continue
                require((sheet.max_row or 0) <= MAX_ROWS + 30, f'{sheet.title} 超过 2000 行，请分批导入', 422)
                # Inherit only real, single-column merged ranges, never arbitrary blanks.
                merged_cells = {}
                xml = ElementTree.fromstring(archive.read(sheet._worksheet_path))
                for node in xml.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}mergeCell'):
                    left, top, right, bottom = range_boundaries(node.attrib['ref'])
                    if left == right and left in columns.values() and index + 2 <= top <= bottom <= MAX_ROWS + 30:
                        for row_number in range(top + 1, bottom + 1):
                            merged_cells[(row_number, left)] = (top, left)
                merged_values = {}
                sheet_rows = []
                raw_rows = sheet.iter_rows(min_row=index + 2, max_col=120)
                cached_rows = cached_book[sheet.title].iter_rows(min_row=index + 2, max_col=120)
                for position, (cells, cached_cells) in enumerate(zip(raw_rows, cached_rows), index + 2):
                    source, warnings, notes, evidence = {}, [], [], {}
                    for field, col in columns.items():
                        if field == 'image_data':
                            continue
                        origin = merged_cells.get((position, col))
                        cell, cached = merged_values.get(origin, (cells[col - 1], cached_cells[col - 1]))
                        if origin is None:
                            merged_values[(position, col)] = (cell, cached)
                        value = text(cell)
                        require(len(value) <= 1000, f'{sheet.title} 第 {position} 行字段过长', 422)
                        coordinate = getattr(cell, 'coordinate', None)
                        label = NUMERIC_LABELS.get(field, FIELDS[field][0])
                        if cell.data_type in {'f', 'e'} or value.startswith('='):
                            saved = formula_cache(field, cached) if cell.data_type == 'f' else None
                            if saved is None:
                                warnings.append(f'{label} 含公式但无可用数字结果，或为错误值；请核对后补填或保留已有资料')
                                value = ''
                            else:
                                evidence[field] = dict(cell=coordinate, formula=value, cached_value=str(cached.value), value=saved)
                                notes.append(f'{label}取自 {coordinate} 已保存的计算结果（最多 8 位小数），请核对原表已重新计算并保存。')
                                value = saved
                        if value:
                            source[field] = value
                            if origin:
                                evidence.setdefault(field, {}).update(cell=coordinate, merged=True)
                                notes.append(f'{label}沿用合并区域 {coordinate} 的资料。')
                    if not source.get('item_no') or not (source.get('name') or source.get('code')):
                        if warnings:
                            issues.append(f'{sheet.title} 第 {position} 行未匹配：货号或物料名称／编码为公式或为空，请核对原表')
                        continue
                    if normalized(source.get('name')) in {'合计', '总计'}:
                        continue
                    require(len(result) < MAX_ROWS, '本次物料行超过 2000 行，请分批导入', 422)
                    row = dict(index=len(result), source_sheet=sheet.title, source_row=position, source=source, warnings=warnings, notes=notes, evidence=evidence)
                    result.append(row)
                    sheet_rows.append(row)
                photos, photo_warnings = sheet_photos(archive, sheet._worksheet_path, xml, columns.get('image_data'),
                                                      {row['source_row'] for row in sheet_rows}, image_budget)
                for row in sheet_rows:
                    if row['source_row'] in photos:
                        row['source']['image_data'] = photos[row['source_row']]
                        row['notes'].append('图片取自本行图片列中的内嵌图片。')
                    row['warnings'].extend(photo_warnings.get(row['source_row'], []))
            require(sheets, '未找到货号及物料名称表头，请使用含“货号”和“物料名称／产品中文名称”的统一资料', 422)
            return dict(sheets=sheets, rows=result, issues=issues)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, '无法读取统一资料，请核对文件和表头映射') from exc
    finally:
        if book:
            book.close()
        if cached_book:
            cached_book.close()


def product_matches(value, code):
    return bool(re.search(r'(?<![\w-])' + re.escape(code) + r'(?![\w-])', value, re.IGNORECASE))


def build_groups(ledger, rows, products):
    groups = {}
    materials = ledger.rows('material', include_deleted=True)
    for source_row in rows:
        source = source_row['source']
        targets = [p for p in products if product_matches(source['item_no'], p.data['code'])]
        if not targets:
            continue
        code = source.get('code', '')
        matches = [m for m in materials if
            value_key(m.data['name']) == value_key(source.get('name')) and
            value_key(m.data.get('specification')) == value_key(source.get('specification'))]
        name_matches = [dict(id=m.id, version=m.version) for m in matches]
        if len(matches) > 1 and code:
            exact_code = [m for m in matches if value_key(m.data['code']) == value_key(code)]
            if exact_code:
                matches = exact_code
        existing = matches[0] if len(matches) == 1 else None
        key = ('existing', existing.id) if existing else ('name', value_key(source.get('name')), value_key(source.get('specification')))
        if key not in groups:
            affected = [dict(id=p.id, code=p.data['code'], customer=p.data['customer']) for p in ledger.rows('product', include_deleted=True)
                        if existing and any(r['material_id'] == existing.id for r in p.data['materials'])]
            groups[key] = dict(index=len(groups), existing=view(existing) if existing else None, sources=[], products=[], affected_products=affected,
                               name_matches=name_matches,
                               blocked='缺少产品中文名称，请补齐原表后重新预览' if not source.get('name') else
                               '相同名称和规格对应多条物料，请补充编码或核对基础资料后重新预览' if len(matches) > 1 else
                               '物料编码已删除，请先恢复原物料；导入不能绕过删除状态' if existing and existing.data.get('deleted') else '')
        group = groups[key]
        group['sources'].append(deepcopy(source_row) | {'product_ids': [p.id for p in targets]})
        for p in targets:
            if not any(item['id'] == p.id for item in group['products']):
                recipe = next((r for r in p.data['materials'] if existing and r['material_id'] == existing.id), None)
                group['products'].append(dict(id=p.id, code=p.data['code'], existing_usage=recipe['usage'] if recipe else None))
    names_by_code = defaultdict(set)
    for row in materials:
        names_by_code[value_key(row.data['code'])].add(row.data['name'])
    for row in rows:
        names_by_code[value_key(row['source'].get('code'))].add(row['source'].get('name', ''))
    for group in groups.values():
        group['code_warnings'] = []
        for code in sorted({r['source'].get('code', '') for r in group['sources']} - {''}):
            names = sorted(names_by_code[value_key(code)] - {''})
            if len(names) > 1:
                group['code_warnings'].append(f'编码 {code} 对应 {len(names)} 种中文名称，当前物料独立维护。')
        if group['existing'] and any(r['source'].get('code') and r['source']['code'] != group['existing']['data']['code'] for r in group['sources']):
            group['code_warnings'].append('中文名称及规格匹配已有物料，来源编码不同；保留已有编码，请核对资料。')
    return list(groups.values())


def preview(db, user, factory, customer, product_ids, filename, content, parsed, refresh_user):
    try:
        lock_mutation(db)
        lock_transaction(db, 'indonesia-materials', factory)
        current = refresh_user()
        require(current.id == user.id, '账号身份已变化，请重新上传核对', 401)
        authorize(current, factory, 'write')
        ledger = Ledger(db, factory)
        ledger.fixed_customer(customer)
        require(0 < len(product_ids) <= 50 and len(set(product_ids)) == len(product_ids), '请选择 1 至 50 个不同货号', 422)
        products = [ledger.get(id, 'product') for id in product_ids]
        require(all(p.data['customer'] == customer for p in products), '所选货号不属于当前固定客人', 422)
        groups = build_groups(ledger, parsed['rows'], products)
        # A fresh version snapshot is needed after a concurrent edit; reuploading must not reuse stale comparisons.
        key = [TYPE, IDENTITY_RULE, FORMAT_RULE, hashlib.sha256(content).hexdigest(), customer, sorted(product_ids), parsed['sheets'],
               [[p.id, p.version] for p in products], [g['name_matches'] for g in groups],
               [[g['existing']['id'], g['existing']['version'], g['affected_products']] for g in groups if g['existing']]]
        preview_key = digest(key)
        previous = next((r for r in ledger.rows('import') if r.data.get('preview_key') == preview_key and r.data['status'] == 'preview'), None)
        if previous:
            return dict(record=view(previous))
        asset = ledger.asset(PurePath(filename).name, content, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        record = ledger.save('import', dict(import_type=TYPE, identity_rule=IDENTITY_RULE, format_rule=FORMAT_RULE, status='preview', filename=asset.filename, asset_id=asset.id,
            customer=customer, products=[dict(id=p.id, version=p.version, code=p.data['code']) for p in products],
            sheets=parsed['sheets'], issues=parsed.get('issues', []), groups=groups, matched_rows=sum(len(g['sources']) for g in groups),
            skipped_rows=len(parsed['rows']) - sum(len(g['sources']) for g in groups), provided_by=user.id, preview_key=preview_key))
        db.commit()
        return dict(record=view(record))
    except Exception:
        db.rollback()
        raise


def confirm(ledger, record, payload, actor):
    require(record.data.get('import_type') == TYPE and record.data['status'] == 'preview', '请选择未确认的统一资料导入记录', 409)
    require(record.data.get('identity_rule') == IDENTITY_RULE, '物料匹配规则已更新，请重新读取最新资料并对比；原草稿保留', 409)
    require(record.data.get('format_rule') == FORMAT_RULE, '模板识别规则已更新，请重新读取最新资料并对比；原草稿保留', 409)
    selection = s.MasterImportSelection.model_validate(payload)
    ledger.fixed_customer(record.data['customer'])
    products = {}
    for prior in record.data['products']:
        product = ledger.get(prior['id'], 'product')
        require(product.version == prior['version'], f"货号 {prior['code']} 已被修改，请重新预览；本次未保存", 409)
        products[product.id] = product
    groups = {g['index']: g for g in record.data['groups']}
    require(len({c.group for c in selection.choices}) == len(selection.choices), '同一物料不能重复选择', 422)
    updates, imported, selected_materials = {}, [], set()
    original_versions = {m.id: m.version for m in ledger.rows('material', include_deleted=True)}
    for choice in selection.choices:
        group = groups.get(choice.group)
        require(group is not None and not group['blocked'], '所选物料无法导入，请核对匹配提示', 422)
        require(any(r['index'] == choice.source_index for r in group['sources']), '来源行不属于此物料', 422)
        existing = group['existing']
        material = ledger.get(existing['id'], 'material') if existing else None
        require(not material or material.version == existing['version'], '物料资料已被修改，请重新预览对比；本次未保存', 409)
        target_ids = {p['id'] for p in group['products']}
        require({r.product_id for r in choice.recipes} == target_ids and len(choice.recipes) == len(target_ids), '货号用量须与本物料匹配货号一致', 422)
        if choice.decision == 'keep':
            require(material is not None and choice.material is None, '保留已有资料不能携带覆盖内容', 422)
        else:
            require(choice.material is not None, '请补齐导入物料资料', 422)
            if material:
                current_refs = {p.id for p in ledger.rows('product', include_deleted=True)
                                if any(r['material_id'] == material.id for r in p.data['materials'])}
                require(current_refs == {p['id'] for p in group['affected_products']}, '共用物料的关联货号已变化，请重新预览影响范围', 409)
            data = choice.material.model_dump(mode='json')
            require(not material or data['code'] == material.data['code'], '不能用导入改变已有物料编码', 422)
            source = group['sources'][0]['source']
            same_identity = all(value_key(data.get(key)) == value_key(source.get(key)) for key in ('name', 'specification'))
            original_peers = {(m['id'], m['version']) for m in group['name_matches'] if not material or m['id'] != material.id} if same_identity else set()
            current_peers = {(m.id, original_versions.get(m.id, m.version) if m.id in selected_materials else m.version)
                             for m in ledger.rows('material', include_deleted=True)
                             if m.id != (material.id if material else None) and
                             all(value_key(m.data.get(key)) == value_key(data.get(key)) for key in ('name', 'specification'))}
            require(current_peers == original_peers, '同名同规格物料已变化或在预览后新增，请重新预览对比', 409)
            data['billing_mode'] = 'automatic'
            material = ledger.save_material(data, material)
        require(material.id not in selected_materials, '同一物料重复匹配，请重新预览核对', 409)
        selected_materials.add(material.id)
        for recipe in choice.recipes:
            data = updates.setdefault(recipe.product_id, deepcopy(products[recipe.product_id].data))
            prior = next((r for r in data['materials'] if r['material_id'] == material.id), None)
            if recipe.decision == 'keep':
                require(prior is not None and recipe.usage is None, '无已有用量可保留，请填写单件用量', 422)
            else:
                require(recipe.usage is not None, '请填写单件用量；不能使用历史送货数量代替', 422)
                entry = dict(material_id=material.id, usage=str(recipe.usage))
                if prior:
                    prior.update(entry)
                else:
                    data['materials'].append(entry)
            require(len(data['materials']) <= 100, '一个货号最多关联 100 种物料', 422)
        imported.append(dict(group=choice.group, material_id=material.id, source_index=choice.source_index,
                             decision=choice.decision, recipes=[r.model_dump(mode='json') for r in choice.recipes]))
    for product_id, data in updates.items():
        if data != products[product_id].data:
            ledger.save('product', data, products[product_id])
    data = deepcopy(record.data)
    data.update(status='confirmed', choices=imported, confirmed_by=actor)
    return ledger.save('import', data, record)
