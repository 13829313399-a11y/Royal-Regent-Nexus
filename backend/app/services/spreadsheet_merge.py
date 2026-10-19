"""Bounded in-memory spreadsheet inspection, explicit merge and audited export.

This service has no storage registry and never opens a business database.
"""
from collections import Counter
import csv
from datetime import date, datetime, time
from io import BytesIO, StringIO
import math
from pathlib import PurePosixPath
import re
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile, ZIP_DEFLATED

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter

from app.schemas.spreadsheet_merge import MergeRequest

MAX_BYTES = 64 * 1024 * 1024
MAX_EXPANDED = 160 * 1024 * 1024
MAX_CELLS = 250000
MAX_SHEETS = 50
BUSINESS_MAX_BYTES = 5 * 1024 * 1024
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def fail(message, status=422):
    raise HTTPException(status, message)


def preflight_xlsx(content):
    """Check ZIP size before any decoder; scan actual values, ignoring style-only cells."""
    bounds = {}
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 5000 or sum(m.file_size for m in members) > MAX_EXPANDED:
                fail('工作簿解压后不能超过 160 MiB 或 5000 个文件', 413)
            names = [m.filename for m in members]
            if len(names) != len(set(names)):
                fail('工作簿包含重复 ZIP 路径')
            total, rectangular, sheet_count = 0, 0, 0
            for member in members:
                if member.flag_bits & 1:
                    fail('不支持加密工作簿')
                path = PurePosixPath(member.filename)
                if path.is_absolute() or '..' in path.parts or '\\' in member.filename:
                    fail('工作簿路径无效')
                if not member.filename.endswith(('.xml', '.rels')):
                    continue
                # Do not allow DTD/entity expansion in OOXML, including shared strings.
                raw = archive.read(member)
                if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():
                    fail('工作簿 XML 含不支持的实体声明')
                # OOXML relationship targets need not use Excel's sheet1.xml naming.
                # Inspect every worksheet XML part, including custom filenames and
                # nested paths, before any reader allocates its effective rectangle.
                if not member.filename.startswith('xl/worksheets/') or not member.filename.endswith('.xml'):
                    continue
                sheet_count += 1
                if sheet_count > MAX_SHEETS:
                    fail('工作簿最多 50 个工作表', 413)
                max_row, max_col = 0, 0
                for _, cell in ET.iterparse(BytesIO(raw), events=('end',)):
                    if cell.tag != NS + 'c':
                        if cell.tag == NS + 'row':
                            cell.clear()
                        continue
                    meaningful = cell.find(NS + 'f') is not None or any(
                        node.text is not None and node.text != '' for node in cell.iter() if node.tag in {NS + 'v', NS + 't'})
                    if meaningful:
                        col, row = coordinate_from_string(cell.attrib['r'])
                        column = column_index_from_string(col)
                        if row > 5001 or column > 256:
                            fail('每页有效区域最多 5000 数据行、256 列（请移除远端数据）', 413)
                        max_row, max_col = max(max_row, row), max(max_col, column)
                        total += 1
                        if total > MAX_CELLS:
                            fail('工作簿实际单元格超过 250000 个', 413)
                    cell.clear()
                bounds[member.filename] = (max_row, max_col)
                rectangular += max_row * max_col
                if rectangular > MAX_CELLS:
                    fail('工作簿有效数据矩形区域合计超过 250000 个单元格，请移除远端稀疏数据', 413)
    except HTTPException:
        raise
    except (BadZipFile, ET.ParseError, KeyError, ValueError, OSError, RuntimeError) as exc:
        raise HTTPException(422, '无法读取 XLSX，请检查文件是否完整') from exc
    return bounds


def checked_worksheet_bounds(bounds, worksheet_path):
    """Every worksheet actually opened by the decoder must have passed preflight.

    Unusual relationship targets outside the supported worksheet tree are refused
    explicitly; absence of a bounds record must never masquerade as an empty sheet.
    """
    key = worksheet_path.lstrip('/')
    if key not in bounds:
        fail('工作表路径未完成安全检查，不支持此工作簿结构；请另存为标准 XLSX 后重试')
    return bounds[key]


def _scalar(value):
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value)


def inspect(content, filename, preset='generic'):
    if len(content) > MAX_BYTES:
        fail('文件不能超过 64 MiB', 413)
    if not content:
        fail('文件为空')
    suffix = PurePosixPath(filename).suffix.lower()
    warnings = []
    if suffix == '.csv':
        if preset != 'generic':
            fail('专用尺码表适配只支持 XLSX')
        text = None
        for encoding in ('utf-8-sig', 'gb18030'):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                pass
        if text is None:
            fail('CSV 编码无法读取，请另存为 UTF-8')
        rows = []
        count = 0
        try:
            for row in csv.reader(StringIO(text)):
                if len(row) > 256 or len(rows) >= 5001:
                    fail('CSV 最多 5000 数据行、256 列', 413)
                if any(len(v) > 32767 for v in row):
                    fail('单元格文本超过 Excel 长度限制')
                count += len(row)
                if count > MAX_CELLS:
                    fail('CSV 单元格超过 250000 个', 413)
                rows.append(row)
        except csv.Error as exc:
            raise HTTPException(422, 'CSV 格式无效') from exc
        return {'filename': filename, 'sheets': [{'name': 'CSV', 'rows': rows}], 'warnings': ['CSV 按文本保留前导零；不会执行公式。']}
    if suffix != '.xlsx':
        fail('仅支持 XLSX 或 CSV，请先另存为这些格式')
    bounds = preflight_xlsx(content)
    if preset == 'carton-packaging-source':
        from app.services.carton_packaging_adapter import adapt_carton_workbook
        return {'filename': filename, **adapt_carton_workbook(content, bounds=bounds)}
    formula_book = cached_book = None
    try:
        formula_book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        cached_book = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        sheets, formula_count, missing_cache, numeric_ids = [], 0, 0, 0
        for sheet in formula_book:
            # openpyxl resolves relationship paths to each worksheet; never trust dimensions.
            row_count, col_count = checked_worksheet_bounds(bounds, sheet._worksheet_path)
            rows = []
            identifiers = set()
            cached = cached_book[sheet.title]
            sheet.reset_dimensions()
            cached.reset_dimensions()
            if not row_count:
                sheets.append({'name': sheet.title, 'rows': []})
                continue
            for row_no, (cells, caches) in enumerate(zip(sheet.iter_rows(max_row=row_count, max_col=col_count), cached.iter_rows(max_row=row_count, max_col=col_count)), 1):
                row = []
                for column, (cell, cache) in enumerate(zip(cells, caches)):
                    if cell.data_type == 'f':
                        formula_count += 1
                        value = None if cache.data_type == 'e' else _scalar(cache.value)
                        missing_cache += value is None
                    else:
                        value = _scalar(cell.value)
                    if row_no <= 20 and isinstance(value, str) and re.search(r'货号|款号|物料编码|item.?no|article.?no|product.?code', value, re.I):
                        identifiers.add(column)
                    if column in identifiers and isinstance(value, (int, float)) and not isinstance(value, bool):
                        numeric_ids += 1
                        value = str(int(value)) if float(value).is_integer() else str(value)
                        if re.fullmatch(r'0{2,}', cell.number_format or '') and re.fullmatch(r'\d+', value):
                            value = value.zfill(len(cell.number_format))
                    row.append(value)
                while row and row[-1] is None:
                    row.pop()
                rows.append(row)
            sheets.append({'name': sheet.title, 'rows': rows})
        if formula_count:
            warnings.append(f'{formula_count} 个公式仅读取原文件缓存，不执行公式；其中 {missing_cache} 个无有效缓存，已保留为空待核。')
        if numeric_ids:
            warnings.append(f'{numeric_ids} 个数字货号已转为文本，纯补零格式已保留；原文件中已丢失的前导零或超过 Excel 精度的数字无法恢复，请核对。')
        warnings.append('仅提取单元格值；图片、格式和合并区域不会带入结果。')
        return {'filename': filename, 'sheets': sheets, 'warnings': warnings}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, '无法读取工作簿数据，请检查文件格式') from exc
    finally:
        if formula_book is not None:
            formula_book.close()
        if cached_book is not None:
            cached_book.close()


def empty(value):
    return value is None or isinstance(value, str) and not value.strip()


def _key(row, keys):
    if any(empty(row[c]) for c in keys):
        return None
    # Typed, exact keys avoid collapsing 001, 1, True and '1'. Whitespace is significant.
    return tuple((type(row[c]).__name__, row[c]) for c in keys)


def _write_cell(sheet, row, column, value):
    if isinstance(value, str) and len(value) > 32767:
        fail('输出单元格文本超过 Excel 长度限制，请缩短备注或列名')
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        # Strings beginning with =,+,-,@ remain literal strings; never emit executable formulas.
        cell.data_type = 's'
        cell.number_format = '@'
    return cell


def _workbook(headers, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = '导入数据'
    for r, values in enumerate([headers, *rows], 1):
        for c, value in enumerate(values, 1):
            _write_cell(sheet, r, c, value)
    sheet.freeze_panes = 'A2'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0F766E')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[1].height = 32
    for column, header in enumerate(headers, 1):
        width = 72 if '备注' in header else 28 if header in {'货号', '产品名称', '配置组'} else 18
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    return book


def _bytes(book):
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class _EmptyDatabase:
    def scalars(self, statement):
        return []


def _batches(payload, rows):
    size = payload.batch_size
    if size is None:
        return [list(range(len(rows)))]
    if payload.profile != 'carton-packaging':
        return [list(range(start, min(start + size, len(rows)))) for start in range(0, len(rows), size)] or [[]]
    from app.services.carton_master_import import HEADERS, norm
    if payload.primary.headers != HEADERS['configurations']:
        return [list(range(start, min(start + size, len(rows)))) for start in range(0, len(rows), size)] or [[]]
    groups = {}
    for index, row in enumerate(rows):
        # Missing identities remain pending independent rows, never an invented shared group.
        group_name = '默认' if empty(row[2]) else str(row[2])
        key = (norm(str(row[0])), norm(group_name)) if not empty(row[0]) else ('missing', index)
        groups.setdefault(key, []).append(index)
    batches, batch = [], []
    for group in groups.values():
        if len(group) > size:
            fail(f'同一货号与配置组有 {len(group)} 行，不能拆开；请将每批行数设为至少 {len(group)}（最多 1000）')
        if batch and len(batch) + len(group) > size:
            batches.append(batch)
            batch = []
        batch.extend(group)
    if batch:
        batches.append(batch)
    return batches or [[]]


def _validate_profile(payload, rows):
    if payload.template_id:
        from app.services import spreadsheet_merge_templates as templates
        templates.validate_selection(payload.template_id, payload.primary.headers, payload.profile)
        if payload.template_id != 'carton-packaging':
            return [f"{templates.get_template(payload.template_id)['label']}：转换完成后仍须进入原业务导入预览；本工具未校验该业务的必填值、分组、厂区现有资料与导入设置。"]
    if payload.profile == 'generic':
        return []
    from app.services import carton_master_import as contract
    if payload.primary.headers != contract.HEADERS['configurations']:
        return ['货号包装必须使用系统模板的完整列名和顺序']
    errors, entries = [], []
    for part, indexes in enumerate(_batches(payload, rows), 1):
        chunk = [rows[index] for index in indexes]
        try:
            parsed, problems = contract.parse(_bytes(_workbook(payload.primary.headers, chunk)), 'merge.xlsx', 'configurations')
            # Remap row references to the full output, avoiding misleading batch-local positions.
            def remap(message):
                return re.sub(r'第 (\d+) 行', lambda m: f'第 {indexes[int(m.group(1)) - 2] + 2} 行' if 2 <= int(m.group(1)) <= len(indexes) + 1 else m.group(0), message)
            errors.extend(remap(message) for message in problems)
            entries.extend((indexes[index - 2] + 2, entry) for index, entry in parsed)
        except HTTPException as exc:
            errors.append(f'第 {part} 批：{exc.detail}')
    planned, _ = contract.plan(_EmptyDatabase(), 'tool-validation', 'configurations', entries, [])
    errors.extend(planned['errors'])
    if len(rows) > 1000 and payload.batch_size is None:
        errors.append('业务导入每批最多 1000 行，请启用分批下载')
    if not rows:
        errors.append('没有可导入的数据')
    return list(dict.fromkeys(errors))


def merge(payload: MergeRequest):
    width = len(payload.primary.headers)
    primary = [list(row) + [None] * (width - len(row)) for row in payload.primary.rows]
    secondary = []
    for original in payload.secondary.rows:
        row = [None] * width
        for mapping in payload.mapping:
            row[mapping.target] = original[mapping.source] if mapping.source < len(original) else None
        if payload.extra_columns_to_remarks:
            extras = [f'{payload.secondary.headers[c]}：{original[c]}' for c in payload.extra_columns_to_remarks if c < len(original) and not empty(original[c])]
            if extras:
                current = row[payload.remarks_column]
                row[payload.remarks_column] = '\n'.join([*([str(current)] if not empty(current) else []), *extras])
                if len(row[payload.remarks_column]) > 32767:
                    fail('拼入备注后的文本超过 Excel 长度限制，请减少附加字段')
        secondary.append(row)
    rows = [r[:] for r in primary]
    def source(table, index):
        data = getattr(payload, table)
        return {'table': table, 'name': data.name, 'row': data.header_row + index + 1}
    provenance = [[source('primary', i)] for i in range(len(primary))]
    conflicts, issues, warnings = [], [], []
    keys = payload.key_columns
    pkeys = [_key(r, keys) for r in primary] if keys else []
    skeys = [_key(r, keys) for r in secondary] if keys else []
    pc, sc = Counter(pkeys), Counter(skeys)
    lookup = {key: index for index, key in enumerate(pkeys) if key is not None and pc[key] == 1}
    filled, added = 0, 0
    if keys:
        for i, key in enumerate(pkeys):
            if key is None or pc[key] > 1 or sc[key] > 1:
                issues.append({'code': 'EMPTY_KEY' if key is None else 'DUPLICATE_KEY', 'message': '主表匹配键为空或重复，保留原行并等待核对', 'rows': [i]})
    for i, incoming in enumerate(secondary):
        key = skeys[i] if keys else None
        ambiguous = bool(keys) and (key is None or sc[key] > 1 or pc[key] > 1)
        target = lookup.get(key) if keys and not ambiguous else None
        if target is None:
            rows.append(incoming[:])
            provenance.append([source('secondary', i)])
            added += 1
            if ambiguous:
                issues.append({'code': 'EMPTY_KEY' if key is None else 'DUPLICATE_KEY', 'message': '副表匹配键为空或重复，未自动合并，已并列保留', 'rows': [len(rows) - 1]})
            continue
        provenance[target].append(source('secondary', i))
        for column, value in enumerate(incoming):
            if empty(value):
                continue
            current = rows[target][column]
            if empty(current):
                rows[target][column] = value
                filled += 1
            elif type(current) is not type(value) or current != value:
                conflict_id = f'p{target}:s{i}:c{column}'
                resolution = payload.resolutions.get(conflict_id)
                conflicts.append({'id': conflict_id, 'row': target, 'column': column, 'primary': current, 'secondary': value, 'resolution': resolution})
                if resolution == 'secondary':
                    rows[target][column] = value
    # Staging warnings remain blocking even when their column is deliberately unmapped.
    # Mapping a plausible scalar cannot erase a source's unresolved business ambiguity.
    source_outputs = {}
    for output_index, references in enumerate(provenance):
        for reference in references:
            source_outputs.setdefault((reference['table'], reference['row']), []).append(output_index)
    for table_name in ('primary', 'secondary'):
        table = getattr(payload, table_name)
        problem_columns = [c for c, header in enumerate(table.headers) if header.strip() == '待核问题']
        for index, original in enumerate(table.rows):
            for column in problem_columns:
                if column < len(original) and not empty(original[column]):
                    source_row = table.header_row + index + 1
                    output_rows = source_outputs.get((table_name, source_row), [])
                    issues.append({'code': 'SOURCE_REVIEW_REQUIRED', 'message': f'{table.name} 第 {source_row} 行：{original[column]}', 'rows': output_rows})
    unknown = set(payload.resolutions) - {c['id'] for c in conflicts}
    if unknown:
        fail('冲突选择已失效，请重新预览后确认')
    unselected = set(range(len(payload.secondary.headers))) - {m.source for m in payload.mapping} - set(payload.extra_columns_to_remarks)
    if unselected:
        warnings.append('未映射副表列未加入数据页，原值保留在说明页：' + '、'.join(payload.secondary.headers[c] for c in sorted(unselected)))
    validation = _validate_profile(payload, rows)
    unresolved = any(c['resolution'] is None for c in conflicts)
    return {'headers': payload.primary.headers, 'rows': rows, 'provenance': provenance, 'conflicts': conflicts,
            'issues': issues, 'warnings': warnings, 'validation_errors': validation,
            'ready_for_import': not unresolved and not issues and not validation,
            'stats': {'primary_rows': len(primary), 'secondary_rows': len(secondary), 'output_rows': len(rows), 'filled_cells': filled, 'added_rows': added}}


def _audited_workbook(payload, result, indexes, first=True):
    if payload.template_id:
        from app.services import spreadsheet_merge_templates as templates
        book = templates.create_workbook(payload.template_id, result['headers'], [result['rows'][index] for index in indexes])
        audit_name = templates.audit_sheet_name(payload.template_id)
        first_data_row = templates.get_template(payload.template_id)['header_row'] + 1
        notes = None
    else:
        book = _workbook(result['headers'], [result['rows'][index] for index in indexes])
        notes = book.create_sheet('说明与示例')
        audit_name = '说明与示例'
        first_data_row = 2
    # Keep all audits in the only non-import sheet accepted by the real business parser.
    records = [
        ['状态', '可交由业务导入再次校验' if result['ready_for_import'] else '待补版：不能直接业务导入，请先处理待核问题'],
        ['边界', '仅转换与核对；不会写入任何业务数据。业务导入仍须自己的权限、厂区与现有资料校验。'],
        ['校验配置', payload.profile], ['本文件数据行数', len(indexes)],
        ['示例（不导入）', '文本货号 00123；先处理冲突和缺项，再进入业务导入预览'],
    ]
    records.extend(['警告', warning] for warning in result['warnings'])
    if first:
        records.extend(['待补', error] for error in result['validation_errors'])
    else:
        records.append(['全部校验问题', f'请参见同一 ZIP 内第 001 批文件的{audit_name}页'])
    selected = set(indexes)
    records.extend(['待核', issue['message'], '预览行号', ','.join(str(r + 2) for r in issue['rows'])] for issue in result['issues'] if selected.intersection(issue['rows']))
    records.extend(['冲突', conflict['id'], payload.primary.headers[conflict['column']], conflict['primary'], conflict['secondary'], conflict['resolution']] for conflict in result['conflicts'] if conflict['row'] in selected)
    for local_index, index in enumerate(indexes):
        for ref in result['provenance'][index]:
            records.append(['输出行来源', local_index + first_data_row, ref['table'], ref['name'], ref['row'], '预览行', index + 2])
    # Retain every raw source value, including unmapped columns and conflicting alternatives.
    if first:
        for table_name in ('primary', 'secondary'):
            table = getattr(payload, table_name)
            records.append(['原表列名', table_name, table.name, *table.headers])
            records.extend(['原表数据', table_name, table.header_row + index + 1, *row] for index, row in enumerate(table.rows))
        for table in payload.audit_tables:
            records.append(['审计附表列名', table.name, *table.headers])
            records.extend(['审计附表数据', table.name, table.header_row + index + 1, *row] for index, row in enumerate(table.rows))
    else:
        records.append(['原表完整值', f'请参见同一 ZIP 内第 001 批文件的{audit_name}页'])
    if payload.template_id:
        templates.append_audit(book, payload.template_id, records)
        return book
    for r, values in enumerate(records, 1):
        for c, value in enumerate(values, 1):
            _write_cell(notes, r, c, value)
    notes.column_dimensions['A'].width = 20
    notes.column_dimensions['B'].width = 90
    notes.freeze_panes = 'A2'
    notes.sheet_view.showGridLines = False
    for column in range(3, notes.max_column + 1):
        notes.column_dimensions[get_column_letter(column)].width = 28
    for row in notes:
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    return book


def export(payload: MergeRequest):
    result = merge(payload)
    if any(c['resolution'] is None for c in result['conflicts']):
        fail('请逐项选择冲突使用主表或副表后再下载', 409)
    batches = _batches(payload, result['rows'])
    contents = [_bytes(_audited_workbook(payload, result, indexes, part == 0)) for part, indexes in enumerate(batches)]
    if payload.profile == 'carton-packaging' and any(len(content) > BUSINESS_MAX_BYTES for content in contents):
        result['ready_for_import'] = False
        result['validation_errors'].append('导出文件包含审计原表后超过业务导入 5 MiB 限制；本文件仅为待补核对版，需精简原表或拆分后重新生成')
        contents = [_bytes(_audited_workbook(payload, result, indexes, part == 0)) for part, indexes in enumerate(batches)]
    status = 'ready' if result['ready_for_import'] else 'pending'
    if payload.batch_size is None:
        return contents[0], f'spreadsheet-merge-{status}.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    output = BytesIO()
    with ZipFile(output, 'w', ZIP_DEFLATED) as archive:
        for part, content in enumerate(contents, 1):
            archive.writestr(f'spreadsheet-merge-{status}-{part:03}.xlsx', content)
    return output.getvalue(), f'spreadsheet-merge-{status}.zip', 'application/zip'
