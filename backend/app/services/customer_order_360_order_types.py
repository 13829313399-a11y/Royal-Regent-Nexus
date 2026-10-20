"""Extend only the 360 export copy's active order-type dropdowns."""
from __future__ import annotations

from copy import copy
from hashlib import sha256
import re

from openpyxl.utils.cell import get_column_letter, quote_sheetname, range_to_tuple
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
from openpyxl.worksheet.datavalidation import DataValidation


ORDER_TYPES = ("正单", "加单", "补单", "改单", "备料单")
SHEETS = ("ITEM表", "接单表", "正单评审表")
_HELPER_MARKER = "华康A360订单类型选项引用_V1"
_SECTION_LABELS = {"已完成", "已完成（按验货日期）", "未完成", "未完成订单", "已走货", "已走货订单"}


class OrderTypeValidationError(ValueError):
    pass


def _fail(sheet, message):
    raise OrderTypeValidationError(f"360 订单类型下拉：{sheet.title} {message}")


def _subtract(rect, cut):
    """Subtract an intersecting rectangle without enumerating its cells."""
    x1, y1 = max(rect.min_col, cut.min_col), max(rect.min_row, cut.min_row)
    x2, y2 = min(rect.max_col, cut.max_col), min(rect.max_row, cut.max_row)
    if x1 > x2 or y1 > y2:
        return [rect]
    pieces = []
    for left, top, right, bottom in (
        (rect.min_col, rect.min_row, rect.max_col, y1 - 1),
        (rect.min_col, y2 + 1, rect.max_col, rect.max_row),
        (rect.min_col, y1, x1 - 1, y2),
        (x2 + 1, y1, rect.max_col, y2),
    ):
        if left <= right and top <= bottom:
            pieces.append(CellRange(min_col=left, min_row=top, max_col=right, max_row=bottom))
    return pieces


def _target_ranges(sheet):
    if re.sub(r"\s+", "", str(sheet["B3"].value or "")) != "订单类型":
        _fail(sheet, "B3 不是订单类型，未修改下拉")
    cells = list(sheet._cells.values())
    markers = [c.row for c in cells if c.row >= 4 and c.data_type != "f" and str(c.value or "").strip() == "取消单"]
    if not markers:
        _fail(sheet, "缺少取消单边界，未猜测可编辑范围")
    boundaries = [(c.row, str(c.value or "").strip()) for c in cells
                  if c.column == 1 and c.row >= 4 and c.data_type != "f"
                  and str(c.value or "").strip() in {"取消单", "已走货", "已走货订单", "已完成", "已完成（按验货日期）"}]
    stop = min([*markers, *(row for row, _ in boundaries)])
    ranges = [CellRange(f"B4:B{stop - 1}")] if stop > 4 else []
    completed = sorted({row for row, label in boundaries if label.startswith("已完成")})
    if len(completed) > 1:
        _fail(sheet, "存在多个已完成分区，未猜测订单类型范围")
    if completed:
        start = completed[0] + 1
        last = max((c.row for c in cells if c.value not in (None, "")), default=start - 1)
        end = min((row - 1 for row, _ in boundaries if row >= start), default=last)
        if start <= end:
            ranges.append(CellRange(f"B{start}:B{end}"))
    excluded = [r for r in sheet.merged_cells.ranges if r.min_col <= 2 <= r.max_col]
    excluded.extend(CellRange(f"B{c.row}") for c in cells if c.row >= 4 and c.data_type != "f" and (
        (c.column == 1 and str(c.value or "").strip() in _SECTION_LABELS)
        or (c.column == 2 and str(c.value or "").strip() == "订单类型")
    ))
    for cut in excluded:
        ranges = [piece for rect in ranges for piece in _subtract(rect, cut)]
    return sorted(ranges, key=lambda r: r.min_row)


def _list_source(workbook, sheet, formula):
    """Resolve literal lists or a single fixed range; never evaluate formulas."""
    original = str(formula or "").strip()
    expression = original.removeprefix("=").strip()
    if expression.startswith('"') and expression.endswith('"'):
        options = expression[1:-1].replace('""', '"').split(",")
        missing = [value for value in ORDER_TYPES if value not in options]
        if not missing:
            return {"formula": original}
        options += missing
        literal = '"' + ",".join(options).replace('"', '""') + '"'
        if len(literal) <= 255:
            return {"formula": literal}
        return {"key": "literal:" + literal, "values": options}

    # Sheet-local names take precedence over workbook names, as they do in Excel.
    definition = sheet.defined_names.get(expression) or workbook.defined_names.get(expression)
    reference = str(definition.attr_text if definition is not None else expression).removeprefix("=")
    if "!" not in reference:
        reference = f"{quote_sheetname(sheet.title)}!{reference}"
    try:
        source_name, bounds = range_to_tuple(reference)
        min_col, min_row, max_col, max_row = bounds
        if "[" in source_name or source_name not in workbook.sheetnames:
            raise ValueError("external or missing sheet")
        if not all(bounds) or (min_col != max_col and min_row != max_row):
            raise ValueError("not a one-dimensional fixed range")
        if (max_row - min_row + 1) * (max_col - min_col + 1) > 1024:
            raise ValueError("unbounded option range")
    except (TypeError, ValueError, KeyError):
        _fail(sheet, "订单类型名单不是可安全扩充的字面列表或固定命名区域，请先核对名单公式")
    source = workbook[source_name]
    source_key = getattr(source._cells.get((1, min_col)), "value", None)
    if (definition is not None and min_col == max_col and min_row == 2
            and getattr(source._cells.get((1, 1)), "value", None) == _HELPER_MARKER
            and isinstance(source_key, str)
            and definition.name.startswith("_HX360_OrderTypes_" + sha256(source_key.encode()).hexdigest()[:12])):
        # This managed list already combined the full canonical set with live
        # source references. Formula cells have no cached values in openpyxl;
        # treating those as absent would wrap the helper again on every export.
        return {"formula": original}
    cells = [(row, column, source._cells.get((row, column)))
             for row in range(min_row, max_row + 1) for column in range(min_col, max_col + 1)]
    known = {str(c.value) for _, _, c in cells if c is not None and c.value is not None and c.data_type != "f"}
    missing = [value for value in ORDER_TYPES if value not in known]
    if not missing:
        return {"formula": original}
    values = []
    for row, column, _ in cells:
        address = f"{quote_sheetname(source.title)}!${get_column_letter(column)}${row}"
        values.append(f'=IF({address}="","",{address})')
    return {"key": "range:" + reference, "values": values + missing}


def _materialize_list(workbook, plan):
    if "formula" in plan:
        return plan["formula"]
    helper = next((s for s in workbook if getattr(s._cells.get((1, 1)), "value", None) == _HELPER_MARKER), None)
    if helper is None:
        helper = workbook.create_sheet("_360订单类型选项")
        helper["A1"] = _HELPER_MARKER
        helper.sheet_state = "veryHidden"
    key = plan["key"]
    column = next((c.column for c in helper[1] if c.column > 1 and c.value == key), None)
    if column is None:
        column = max(2, helper.max_column + 1)
        helper.cell(1, column, key)
        for row, value in enumerate(plan["values"], start=2):
            helper.cell(row, column, value)
    letter = get_column_letter(column)
    target = f"{quote_sheetname(helper.title)}!${letter}$2:${letter}${len(plan['values']) + 1}"
    base_name = "_HX360_OrderTypes_" + sha256(key.encode()).hexdigest()[:12]
    name, number = base_name, 1
    while name in workbook.defined_names and workbook.defined_names[name].attr_text != target:
        number += 1
        name = f"{base_name}_{number}"
    if name not in workbook.defined_names:
        workbook.defined_names.add(DefinedName(name, attr_text=target))
    return "=" + name


def ensure_360_order_type_dropdowns(workbook, *, factory_id: str, customer_code: str):
    """Call after insertion/organization and before saving the export copy.

    Values and formulas in all business cells stay untouched. Existing dropdown
    rules are split so titles, merged cells and cancelled/shipped regions retain
    their original rules. A fixed named list gets an independent live reference
    list; its original name, source cells and unrelated validation users stay intact.
    """
    report = {"applied": False, "sheets": {}, "required_options": list(ORDER_TYPES)}
    if factory_id != "huakang-a" or customer_code != "360":
        return report
    plans = []
    # Preflight every sheet before mutating validation rules or adding helpers.
    for name in SHEETS:
        if name not in workbook:
            raise OrderTypeValidationError(f"360 订单类型下拉：缺少 {name}")
        sheet = workbook[name]
        targets = _target_ranges(sheet)
        uncovered = list(targets)
        rules = []
        fallback = None
        for validation in sheet.data_validations.dataValidation:
            affected = []
            for rect in validation.sqref.ranges:
                for target in targets:
                    if rect.min_col <= 2 <= rect.max_col:
                        lo, hi = max(rect.min_row, target.min_row), min(rect.max_row, target.max_row)
                        if lo <= hi:
                            affected.append(CellRange(f"B{lo}:B{hi}"))
            if not affected:
                continue
            if validation.type != "list":
                _fail(sheet, "订单类型存在非列表验证，未覆盖原规则")
            source = _list_source(workbook, sheet, validation.formula1)
            fallback = (validation, source)
            remaining = list(validation.sqref.ranges)
            for target in targets:
                remaining = [piece for rect in remaining for piece in _subtract(rect, target)]
            for cut in affected:
                before = sum(r.max_row - r.min_row + 1 for r in uncovered)
                uncovered = [piece for rect in uncovered for piece in _subtract(rect, cut)]
                if before - sum(r.max_row - r.min_row + 1 for r in uncovered) != cut.max_row - cut.min_row + 1:
                    _fail(sheet, "订单类型存在重叠验证，请先核对重复规则")
            rules.append((validation, remaining, affected, source))
        if uncovered:
            if fallback is None:
                validation = DataValidation(type="list", allow_blank=True)
                source = {"formula": '"' + ",".join(ORDER_TYPES) + '"'}
            else:
                validation, source = fallback
            rules.append((None, [], uncovered, source | {"prototype": validation}))
        plans.append((sheet, targets, rules))
    for sheet, targets, rules in plans:
        for original, remaining, affected, source in rules:
            replacement = copy(original or source["prototype"])
            replacement.formula1 = _materialize_list(workbook, source)
            replacement.sqref = MultiCellRange(affected)
            replacement.showDropDown = False
            if original is not None:
                if remaining:
                    original.sqref = MultiCellRange(remaining)
                else:
                    sheet.data_validations.dataValidation.remove(original)
            sheet.add_data_validation(replacement)
        report["sheets"][sheet.title] = [str(r) for r in targets]
    report["applied"] = True
    return report
