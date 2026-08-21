from __future__ import annotations

import hashlib
import json
import re
from datetime import date, datetime
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from app.schemas.ai.workbook import (
    AIWorkbookHeaderCandidate,
    AIWorkbookHeaderCell,
    AIWorkbookRedactedSample,
    AIWorkbookSemanticSnapshot,
    AIWorkbookSheetSnapshot,
    AIWorkbookSourceLineage,
)
from app.services.injection_scheduling_profiles import normalize_header

INSPECTOR_VERSION = "workbook-semantic-snapshot-v1"
MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
MAX_ZIP_ENTRIES = 5_000
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100
MAX_SHEETS = 50
MAX_ROWS = 100_000
MAX_COLUMNS = 500
MAX_MERGED_REGIONS = 2_000
MAX_FORMULA_CELLS = 10_000
MAX_STYLES = 10_000
MAX_DRAWINGS = 500
_FORMULA_STRING = re.compile(r'"(?:[^"\\]|\\.)*"')
_FORMULA_NUMBER = re.compile(r"(?<![A-Za-z_])\d+(?:\.\d+)?")


class WorkbookInspectionError(ValueError):
    def __init__(self, code: str, message: str, *, status_code: int = 422) -> None:
        super().__init__(message)
        self.code = code
        self.public_message = message
        self.status_code = status_code


def _validate_source(
    source_file_name: str,
    content: bytes,
    *,
    relaxed_limits: bool = False,
) -> tuple[str, str]:
    normalized_name = source_file_name.strip()
    portable = normalized_name.replace("\\", "/")
    if (
        not normalized_name
        or PurePosixPath(portable).name != portable
        or ".." in PurePosixPath(portable).parts
    ):
        raise WorkbookInspectionError("WORKBOOK_UNSAFE_PATH", "文件名或路径不安全")
    suffix = PurePosixPath(portable).suffix.lower()
    if not relaxed_limits and suffix not in {".xlsx", ".xlsm"}:
        raise WorkbookInspectionError(
            "WORKBOOK_UNSUPPORTED_FORMAT",
            "只支持 .xlsx/.xlsm；旧 .xls 请先另存为 xlsx",
        )
    if not content:
        raise WorkbookInspectionError("WORKBOOK_EMPTY", "工作簿内容为空")
    if not relaxed_limits and len(content) > MAX_WORKBOOK_BYTES:
        raise WorkbookInspectionError(
            "WORKBOOK_TOO_LARGE",
            f"工作簿超过 {MAX_WORKBOOK_BYTES // 1024 // 1024} MB 安全上限",
            status_code=413,
        )
    if not content.startswith(b"PK"):
        raise WorkbookInspectionError("WORKBOOK_INVALID_OOXML", "文件不是有效的 OOXML 工作簿")
    return normalized_name, "XLSM" if suffix == ".xlsm" else "XLSX"


def _validate_zip(content: bytes, *, relaxed_limits: bool = False) -> tuple[bool, bool]:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if not relaxed_limits and len(entries) > MAX_ZIP_ENTRIES:
                raise WorkbookInspectionError(
                    "WORKBOOK_TOO_COMPLEX", "工作簿压缩条目过多，已阻断"
                )
            total_uncompressed = 0
            for entry in entries:
                portable = entry.filename.replace("\\", "/")
                path = PurePosixPath(portable)
                if path.is_absolute() or ".." in path.parts:
                    raise WorkbookInspectionError(
                        "WORKBOOK_UNSAFE_ZIP_PATH", "工作簿包含不安全的内部路径"
                    )
                if entry.flag_bits & 0x1:
                    raise WorkbookInspectionError(
                        "WORKBOOK_ENCRYPTED", "不支持加密工作簿"
                    )
                total_uncompressed += entry.file_size
                if not relaxed_limits and total_uncompressed > MAX_UNCOMPRESSED_BYTES:
                    raise WorkbookInspectionError(
                        "WORKBOOK_ZIP_BOMB", "工作簿解压后超过安全上限，已阻断"
                    )
                ratio = entry.file_size / max(entry.compress_size, 1)
                if (
                    not relaxed_limits
                    and entry.file_size > 1_000_000
                    and ratio > MAX_COMPRESSION_RATIO
                ):
                    raise WorkbookInspectionError(
                        "WORKBOOK_ZIP_BOMB", "工作簿压缩比异常，已阻断"
                    )
            names = {entry.filename.replace("\\", "/").lower() for entry in entries}
            return (
                "xl/vbaproject.bin" in names,
                any(name.startswith("xl/drawings/") for name in names),
            )
    except BadZipFile as exc:
        raise WorkbookInspectionError(
            "WORKBOOK_INVALID_OOXML", "文件不是有效的 OOXML 工作簿"
        ) from exc


def _redacted_value(value: object, number_format: str = "") -> tuple[str, str, bool]:
    if value is None:
        return "BLANK", "", False
    if isinstance(value, bool):
        return "BOOLEAN", "<BOOLEAN>", False
    if isinstance(value, (datetime, date)):
        return "DATE", f"<DATE {value:%Y-%m}>", False
    if isinstance(value, (int, float)):
        absolute = abs(float(value))
        digits = len(str(int(absolute))) if absolute >= 1 else 1
        zero_format = bool(re.fullmatch(r"0+", number_format or ""))
        leading_zero = zero_format and len(number_format) > digits
        kind = "integer" if float(value).is_integer() else "decimal"
        return "NUMBER", f"<NUMBER {kind} digits={digits}>", leading_zero
    text = str(value)
    if text.startswith("="):
        shape = _FORMULA_STRING.sub('"…"', text)
        shape = _FORMULA_NUMBER.sub("#", shape)[:200]
        return "FORMULA", "<FORMULA>", False
    leading_zero = len(text) > 1 and text.isdigit() and text.startswith("0")
    label = "DIGITS" if text.isdigit() else "TEXT"
    suffix = " leading_zero=yes" if leading_zero else ""
    return "TEXT", f"<{label} len={len(text)}{suffix}>", leading_zero


def _header_candidates(sheet) -> list[AIWorkbookHeaderCandidate]:
    ranked: list[tuple[float, int, list[AIWorkbookHeaderCell]]] = []
    scan_rows = min(sheet.max_row or 1, 20)
    scan_columns = min(sheet.max_column or 1, 80)
    for row_index in range(1, scan_rows + 1):
        cells: list[AIWorkbookHeaderCell] = []
        nonempty = 0
        text_count = 0
        for column_index in range(1, scan_columns + 1):
            value = sheet.cell(row_index, column_index).value
            if value is None or str(value).strip() == "":
                continue
            nonempty += 1
            if isinstance(value, str) and not value.startswith("="):
                header = value.strip()[:160]
                normalized = normalize_header(header)[:160]
                if normalized:
                    text_count += 1
                    if len(cells) < 40:
                        cells.append(
                            AIWorkbookHeaderCell(
                                column=get_column_letter(column_index),
                                header_text=header,
                                normalized_header=normalized,
                            )
                        )
        if nonempty >= 2 and cells:
            density = nonempty / max(scan_columns, 1)
            text_ratio = text_count / nonempty
            position_bonus = max(0.0, (21 - row_index) / 100)
            confidence = min(1.0, 0.45 * text_ratio + 0.4 * min(density * 4, 1) + position_bonus)
            ranked.append((confidence, row_index, cells))
    ranked.sort(key=lambda item: (-item[0], item[1]))
    return [
        AIWorkbookHeaderCandidate(row=row, confidence=round(score, 4), cells=cells)
        for score, row, cells in ranked[:1]
    ]


def _samples(sheet, display_sheet, headers: list[AIWorkbookHeaderCandidate]) -> list[AIWorkbookRedactedSample]:
    if not headers:
        return []
    primary = headers[0]
    by_column = {cell.column: cell.header_text for cell in primary.cells}
    samples: list[AIWorkbookRedactedSample] = []
    for row_index in range(primary.row + 1, min(primary.row + 3, sheet.max_row) + 1):
        for header in primary.cells[:12]:
            cell = sheet[f"{header.column}{row_index}"]
            kind, raw_sample, leading_zero = _redacted_value(
                cell.value, cell.number_format
            )
            if kind == "BLANK":
                continue
            display_value = display_sheet[cell.coordinate].value
            _, display_sample, display_leading_zero = _redacted_value(
                display_value, cell.number_format
            )
            formula_shape = ""
            if kind == "FORMULA":
                formula_shape = _FORMULA_STRING.sub('"…"', str(cell.value))
                formula_shape = _FORMULA_NUMBER.sub("#", formula_shape)[:200]
            samples.append(
                AIWorkbookRedactedSample(
                    cell_ref=cell.coordinate,
                    header_text=by_column.get(header.column, ""),
                    value_kind=kind,
                    raw_sample=raw_sample,
                    display_sample=display_sample,
                    formula_shape=formula_shape,
                    leading_zero=leading_zero or display_leading_zero,
                )
            )
    return samples[:36]


def inspect_workbook(
    *,
    source_file_name: str,
    content: bytes,
    factory_id: str,
    relaxed_limits: bool = False,
) -> AIWorkbookSemanticSnapshot:
    source_file_name, detected_format = _validate_source(
        source_file_name,
        content,
        relaxed_limits=relaxed_limits,
    )
    has_macros, zip_has_drawings = _validate_zip(
        content,
        relaxed_limits=relaxed_limits,
    )
    keep_vba = detected_format == "XLSM"
    try:
        workbook = load_workbook(BytesIO(content), data_only=False, keep_vba=keep_vba)
        display_workbook = load_workbook(BytesIO(content), data_only=True, keep_vba=keep_vba)
    except (InvalidFileException, OSError, ValueError, BadZipFile) as exc:
        raise WorkbookInspectionError(
            "WORKBOOK_INVALID_OOXML", "无法安全读取该工作簿"
        ) from exc
    try:
        if not workbook.worksheets or (
            not relaxed_limits and len(workbook.worksheets) > MAX_SHEETS
        ):
            raise WorkbookInspectionError(
                "WORKBOOK_SHEET_LIMIT", f"Sheet 数必须在 1 到 {MAX_SHEETS} 之间"
            )
        style_count = len(getattr(workbook, "_cell_styles", ()))
        if not relaxed_limits and style_count > MAX_STYLES:
            raise WorkbookInspectionError(
                "WORKBOOK_TOO_COMPLEX", "样式数量超过安全上限，已阻断"
            )
        sheets: list[AIWorkbookSheetSnapshot] = []
        total_formulas = 0
        total_drawings = 0
        for sheet, display_sheet in zip(
            workbook.worksheets, display_workbook.worksheets, strict=True
        ):
            max_row = max(sheet.max_row or 1, 1)
            max_column = max(sheet.max_column or 1, 1)
            if not relaxed_limits and (max_row > MAX_ROWS or max_column > MAX_COLUMNS):
                raise WorkbookInspectionError(
                    "WORKBOOK_DIMENSION_LIMIT",
                    f"Sheet {sheet.title} 超过 {MAX_ROWS} 行或 {MAX_COLUMNS} 列安全上限",
                )
            merged = [str(item) for item in sheet.merged_cells.ranges]
            if not relaxed_limits and len(merged) > MAX_MERGED_REGIONS:
                raise WorkbookInspectionError(
                    "WORKBOOK_TOO_COMPLEX",
                    f"Sheet {sheet.title} 合并区域过多，已阻断",
                )
            formula_count = sum(
                1
                for cell in getattr(sheet, "_cells", {}).values()
                if isinstance(cell.value, str) and cell.value.startswith("=")
            )
            total_formulas += formula_count
            if not relaxed_limits and total_formulas > MAX_FORMULA_CELLS:
                raise WorkbookInspectionError(
                    "WORKBOOK_TOO_COMPLEX", "公式单元格超过安全上限，已阻断"
                )
            total_drawings += len(getattr(sheet, "_images", ())) + len(
                getattr(sheet, "_charts", ())
            )
            if not relaxed_limits and total_drawings > MAX_DRAWINGS:
                raise WorkbookInspectionError(
                    "WORKBOOK_TOO_COMPLEX", "图片或图表数量超过安全上限，已阻断"
                )
            candidates = _header_candidates(sheet)
            sheets.append(
                AIWorkbookSheetSnapshot(
                    name=sheet.title,
                    state=sheet.sheet_state,
                    max_row=max_row,
                    max_column=max_column,
                    dimension=sheet.calculate_dimension(),
                    merged_region_count=len(merged),
                    merged_regions=merged[:100],
                    hidden_row_count=sum(
                        bool(item.hidden) for item in sheet.row_dimensions.values()
                    ),
                    hidden_column_count=sum(
                        bool(item.hidden) for item in sheet.column_dimensions.values()
                    ),
                    formula_cell_count=formula_count,
                    candidate_headers=candidates,
                    redacted_samples=_samples(sheet, display_sheet, candidates),
                )
            )
        source_sha256 = hashlib.sha256(content).hexdigest()
        payload = {
            "factory_id": factory_id,
            "source_sha256": source_sha256,
            "sheets": [item.model_dump(mode="json") for item in sheets],
        }
        snapshot_sha256 = hashlib.sha256(
            json.dumps(
                payload, ensure_ascii=False, separators=(",", ":"), sort_keys=True
            ).encode("utf-8")
        ).hexdigest()
        warnings = [
            "样例值已脱敏；前导零、公式原文形态、缓存显示值和日期类型分别记录。",
            "此快照不创建导入批次、订单、任务或 Profile，也不修改源文件。",
        ]
        return AIWorkbookSemanticSnapshot(
            factory_id=factory_id,
            source_lineage=AIWorkbookSourceLineage(
                source_file_name=source_file_name,
                source_sha256=source_sha256,
                source_size_bytes=len(content),
                detected_format=detected_format,
                inspector_version=INSPECTOR_VERSION,
            ),
            sheet_count=len(sheets),
            formula_cell_count=total_formulas,
            hidden_sheet_count=sum(item.state != "visible" for item in sheets),
            has_macros=has_macros,
            has_drawings=zip_has_drawings or total_drawings > 0,
            snapshot_sha256=snapshot_sha256,
            sheets=sheets,
            warnings=warnings,
        )
    finally:
        workbook.close()
        display_workbook.close()
