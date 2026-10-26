"""Read-only legacy stock inspection. No ledger, receipt or master writes."""
import io
import re
import zipfile
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import PureWindowsPath

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.services import fabric_master, warehouse_locations

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
MAX_PROBLEMS = 500
MAX_DUPLICATE_GROUPS = 200
MAX_DUPLICATE_ROWS = 1000
FIELDS = {
    'supplier': ('供应商',), 'category': ('布料/辅料类别', '布料', '辅料类别', '物料类别'),
    'received_on': ('入库日期',), 'location': ('位置', '仓位'),
    'new_item_code': ('新物料编号',), 'item_code': ('物料编号',), 'item_name': ('物料名称规格', '物料名称/规格'),
    'unit': ('单位',), 'customer': ('客名', '客户名称'), 'contract': ('合同',), 'production_no': ('放产编号',),
    'style_no': ('货号',), 'price_note': ('备注单价',), 'unit_price': ('单价',),
    'current_quantity': ('本月结存',), 'previous_quantity': ('上月结存',),
    'paused': ('是否暂停',), 'warehouse_analysis': ('仓库分析',), 'impairment_2026': ('2026报减值',),
}
SUGGESTIONS = {
    'MISSING_QTY': '补齐所选结存列；公式无缓存时先在原表重新计算并保存。',
    'INVALID_QTY': '核对原数量，须为非负、有限的数字，不自动修正或取绝对值。',
    'MISSING_UNIT': '补齐原单位，本次不换算单位。',
    'UNMATCHED_LOCATION': '核对已启用布仓位的完整标签或“仓库／编码”，不按简称猜测。',
    'AMBIGUOUS_LOCATION': '多个仓位匹配，请人工核对正式仓位。',
    'UNMATCHED_MATERIAL': '核对新物料编号、物料编号及已登记的旧编码。',
    'AMBIGUOUS_MATERIAL': '编号或旧编码对应多个物料，请人工核对，不自动合并。',
    'MATERIAL_NOT_ACTIVE': '匹配的物料尚未启用，请在基础资料人工核对。',
    'UNMATCHED_SUPPLIER': '核对供应商编码或名称，不按相似名称匹配。',
    'AMBIGUOUS_SUPPLIER': '供应商编码或名称对应多个记录，请人工核对。',
    'SUPPLIER_NOT_ACTIVE': '匹配的供应商尚未启用，请在基础资料人工核对。',
    'DUPLICATE_ROW': '本行与前一记录的预检字段相同，请核对原表；预检不自动删除或合并。',
    'UNRESOLVED_FORMULA': '原字段公式没有缓存值，请在原表计算保存后重新预检。',
}


def cell_text(cell):
    value = cell.value
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = Decimal(str(value))
        if number.is_finite() and number == number.to_integral_value():
            digits = str(int(number))
            # Preserve Excel's explicit zero padding; never infer absent zeros.
            if re.fullmatch(r'0+', cell.number_format or ''):
                return digits.zfill(len(cell.number_format))
            return digits
    return str(value).strip()


def parse(filename, content):
    if len(content) > MAX_BYTES:
        raise HTTPException(413, '旧库存表不能超过 10 MiB')
    if not filename.lower().endswith('.xlsx') or not content:
        raise HTTPException(422, '请选择有效的 .xlsx 旧库存表')
    books = []
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(part.file_size for part in archive.infolist()) > 200 * 1024 * 1024 or len(archive.infolist()) > 1000:
                raise HTTPException(422, '工作簿解压后过大，请精简后再预检')
        for data_only in (False, True):
            books.append(load_workbook(io.BytesIO(content), read_only=True, data_only=data_only, keep_links=False))
        if '物料' not in books[0].sheetnames:
            raise HTTPException(422, '缺少“物料”工作表')
        sheet, cached = (book['物料'] for book in books)
        columns, labels, header_row = {}, {}, 0
        aliases = {label: key for key, names in FIELDS.items() for label in names}
        for number, cells in enumerate(sheet.iter_rows(max_row=50), 1):
            found = defaultdict(list)
            for index, cell in enumerate(cells):
                label = cell_text(cell)
                if label in aliases:
                    found[aliases[label]].append((index, label))
            if 'item_name' not in found or not ({'new_item_code', 'item_code'} & found.keys()):
                continue
            if any(len(matches) > 1 for matches in found.values()):
                raise HTTPException(422, '预检字段表头重复，请保留每个字段一列')
            columns = {key: matches[0][0] for key, matches in found.items()}
            labels = {key: matches[0][1] for key, matches in found.items()}
            header_row = number
            break
        if not header_row:
            raise HTTPException(422, '前50行未找到物料编号和物料名称规格表头')
        missing = [FIELDS[key][0] for key in ('supplier', 'location', 'unit') if key not in columns]
        quantity_key = 'current_quantity' if 'current_quantity' in columns else 'previous_quantity'
        if quantity_key not in columns:
            missing.append('本月结存或上月结存')
        if missing:
            raise HTTPException(422, '缺少预检表头：' + '、'.join(missing))
        if (sheet.max_row or 0) - header_row > MAX_ROWS:
            raise HTTPException(422, '表头后超过 5000 行，请分批预检并清理范围外格式')
        width = max(columns.values()) + 1
        result = []
        rows = sheet.iter_rows(min_row=header_row + 1, max_row=header_row + MAX_ROWS + 1, max_col=width)
        values = cached.iter_rows(min_row=header_row + 1, max_row=header_row + MAX_ROWS + 1, max_col=width)
        for row_number, (raw_cells, cached_cells) in enumerate(zip(rows, values), header_row + 1):
            if not any(raw_cells[index].value not in (None, '') for index in columns.values()):
                continue
            if row_number > header_row + MAX_ROWS:
                raise HTTPException(422, '表头后超过 5000 行，请分批预检')
            raw = {key: cell_text(raw_cells[index]) for key, index in columns.items()}
            fields = {key: cell_text(cached_cells[index]) for key, index in columns.items()}
            formulas = [key for key, index in columns.items() if raw_cells[index].data_type == 'f']
            unresolved = [key for key in formulas if not fields[key] and key != quantity_key]
            result.append({'row_number': row_number, 'raw': raw, 'values': fields, 'formula_fields': formulas,
                           'unresolved_formulas': unresolved})
        return {'sheet': '物料', 'header_row': header_row, 'sheet_rows': sheet.max_row or 0,
                'field_mapping': {key: {'header': labels[key], 'column': get_column_letter(index + 1)} for key, index in columns.items()},
                'quantity_column': labels[quantity_key], 'quantity_key': quantity_key, 'rows': result}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, '无法读取旧库存表，文件损坏或格式不受支持') from exc
    finally:
        for book in books:
            book.close()


def quantity(value):
    if not value:
        return None, 'MISSING_QTY'
    try:
        number = Decimal(value)
        if not number.is_finite() or number < 0 or len(value) > 128 or abs(number.adjusted()) > 100 or number.as_tuple().exponent < -100:
            raise InvalidOperation()
        return number, None
    except InvalidOperation:
        return None, 'INVALID_QTY'


def normalize_unit(raw):
    value = raw.strip()
    if value in ('', '0', '-', '—', '/', '无'):
        return '', False, '空白或占位值，按缺单位处理，不计入数量合计。'
    if value in ('千克(转算专用)', 'KG（转算专用）'):
        return 'KG（转算专用）', True, '转算专用，单独展示，不计入实物数量合计；不换算数量。'
    if value.lower() == 'pcs':
        return 'PCS', False, 'PCS 大小写别名只归并显示，不换算数量。'
    if value.lower() == 'kg' or value in ('千克', '公斤'):
        return 'KG', False, 'KG、千克、公斤只归并显示，不换算数量。'
    return value, False, '保留原单位，不与其他单位合并，不换算数量。'


def add_total(totals, unit, amount):
    with localcontext() as context:
        context.prec = 256
        totals[unit] += amount


def total_strings(totals):
    return {unit: format(total, 'f') for unit, total in sorted(totals.items())}


def inspect_rows(parsed, masters, locations):
    indexes = {name: defaultdict(set) for name in ('material', 'supplier', 'location')}
    status = {}
    for record in masters:
        family = {'MATERIAL': 'material', 'SUPPLIER': 'supplier'}.get(record['kind'])
        if not family:
            continue
        status[record['id']] = record['status']
        keys = [record['code'], record['data'].get('old_code', '')] if family == 'material' else [record['code'], record['name']]
        for key in keys:
            if key:
                indexes[family][str(key).strip()].add(record['id'])
    for location in locations:
        if location['status'] != 'ACTIVE':
            continue
        for key in {location['label'], f"{location['warehouse']}／{location['code']}", f"{location['warehouse']}/{location['code']}"}:
            indexes['location'][key.strip()].add(location['id'])
    summary = dict(total_rows=len(parsed['rows']), valid_rows=0, zero_rows=0, missing_qty=0, invalid_qty=0,
                   missing_unit=0, duplicate_rows=0, duplicate_groups=0, conversion_only_rows=0)
    matching = {family: dict(matched=0, unmatched=0, ambiguous=0) for family in indexes}
    totals, unique_totals, matched_unique_totals, conversion_totals = (defaultdict(Decimal) for _ in range(4))
    counts, seen, problems, normalized_rows, aliases = defaultdict(int), {}, [], [], {}
    for row in parsed['rows']:
        fields, issues, matches = row['values'], [], {}
        unit, conversion_only, unit_basis = normalize_unit(fields.get('unit', ''))
        raw_unit = row['raw'].get('unit', '')
        alias_key = (raw_unit, unit, conversion_only)
        if alias_key not in aliases:
            aliases[alias_key] = dict(raw=raw_unit, canonical=unit, rows=0, conversion_only=conversion_only)
        aliases[alias_key]['rows'] += 1
        if conversion_only:
            summary['conversion_only_rows'] += 1
        amount, qty_issue = quantity(fields.get(parsed['quantity_key'], ''))
        if qty_issue:
            issues.append(qty_issue)
            summary['missing_qty' if qty_issue == 'MISSING_QTY' else 'invalid_qty'] += 1
        elif amount == 0:
            summary['zero_rows'] += 1
        if not unit:
            issues.append('MISSING_UNIT'); summary['missing_unit'] += 1
        elif amount is not None:
            add_total(conversion_totals if conversion_only else totals, unit, amount)
        if row['unresolved_formulas']:
            issues.append('UNRESOLVED_FORMULA')
        for family, keys in {'material': [fields.get('new_item_code'), fields.get('item_code')],
                             'supplier': [fields.get('supplier')], 'location': [fields.get('location')]}.items():
            identifiers = set().union(*(indexes[family].get(key, set()) for key in keys if key))
            matched = 'unmatched' if not identifiers else 'ambiguous' if len(identifiers) > 1 else 'matched'
            matching[family][matched] += 1
            matches[family] = sorted(identifiers)
            if matched != 'matched':
                issues.append(f'{matched.upper()}_{family.upper()}')
            elif family != 'location' and status[next(iter(identifiers))] != 'ACTIVE':
                issues.append(f'{family.upper()}_NOT_ACTIVE')
        # Keep the original field identity, including the original unit spelling.
        # Unit display aliases must not create new duplicate groups.
        identity = tuple(fields.get(key, '') for key in FIELDS)
        group = seen.setdefault(identity, [])
        duplicate_of = group[0]['row_number'] if group else None
        if duplicate_of is not None:
            issues.append('DUPLICATE_ROW'); summary['duplicate_rows'] += 1
        row_units = dict(raw_unit=raw_unit, unit=unit, unit_basis=unit_basis, conversion_only=conversion_only)
        normalized = dict(row_number=row['row_number'], quantity=format(amount, 'f') if amount is not None else None,
                          raw_quantity=fields.get(parsed['quantity_key'], ''), duplicate_of=duplicate_of, **row_units,
                          **{key: fields.get(key, '') for key in ('new_item_code', 'item_code', 'item_name', 'location', 'supplier')})
        group.append(normalized)
        normalized_rows.append(normalized)
        if duplicate_of is None and unit and amount is not None and not conversion_only:
            add_total(unique_totals, unit, amount)
            if all(len(ids) == 1 and (family == 'location' or status[ids[0]] == 'ACTIVE') for family, ids in matches.items()):
                add_total(matched_unique_totals, unit, amount)
        for code in issues:
            counts[code] += 1
        if not issues:
            summary['valid_rows'] += 1
        elif len(problems) < MAX_PROBLEMS:
            problems.append({**row, 'quantity': format(amount, 'f') if amount is not None else None,
                             'issues': issues, 'suggestions': [SUGGESTIONS[code] for code in issues],
                             'matches': matches, 'duplicate_of': duplicate_of, **row_units})
    groups = [rows for rows in seen.values() if len(rows) > 1]
    summary['duplicate_groups'] = len(groups)
    duplicate_groups, remaining_rows = [], MAX_DUPLICATE_ROWS
    for rows in groups[:MAX_DUPLICATE_GROUPS]:
        if remaining_rows == 0:
            break
        visible = rows[:remaining_rows]
        duplicate_groups.append(dict(key=f"row-{rows[0]['row_number']}", rows=visible, count=len(rows), rows_truncated=len(visible) < len(rows)))
        remaining_rows -= len(visible)
    duplicate_groups_truncated = len(duplicate_groups) < len(groups) or any(group['rows_truncated'] for group in duplicate_groups)
    problem_count = summary['total_rows'] - summary['valid_rows']
    return {key: value for key, value in parsed.items() if key != 'rows'} | {
        'summary': summary, 'matching': matching, 'rows': normalized_rows,
        'unit_aliases': list(aliases.values()), 'conversion_totals': total_strings(conversion_totals),
        'totals': total_strings(totals), 'unique_totals': total_strings(unique_totals),
        'matched_unique_totals': total_strings({unit: matched_unique_totals[unit] for unit in totals}),
        'duplicate_groups': duplicate_groups, 'duplicate_groups_truncated': duplicate_groups_truncated,
        'issue_counts': dict(counts), 'problems': problems, 'problem_count': problem_count,
        'problems_truncated': problem_count > len(problems), 'posted': False, 'read_only': True,
        'quantity_basis': '优先选择本月结存列；仅缺少该列时使用上月结存。公式只读文件缓存，不计算。',
        'totals_basis': '原表合计按归并后的实物单位累计非负有效数量，含重复及未匹配行，不能直接作为系统库存；不换算数量。',
        'unique_totals_basis': '按原始字段判断重复，每组仅计首行；只作排除后续重复行的测算，未合并或删除任何原表行。',
        'matched_unique_totals_basis': '在非重复合计中，仅累计物料、仓位、供应商均唯一匹配且已启用，数量及单位有效的行；不代表允许入账。',
        'conversion_totals_basis': '转算专用行按归并后的单位单独累计，含重复及未匹配行；不计入上述三类实物合计。',
    }


def preview(db, filename, content):
    parsed = parse(filename, content)
    with db.no_autoflush:
        fabric_master.ensure_schema(db)
        masters = [fabric_master.record_data(record) for record in fabric_master.records(db)]
        locations = warehouse_locations.catalog(db, 'fabric')
    return {'filename': PureWindowsPath(filename).name, **inspect_rows(parsed, masters, locations)}
