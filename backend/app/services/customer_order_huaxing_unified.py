"""Huaxing September unified schedules; legacy PO engines remain input-only."""
from __future__ import annotations

from collections import Counter
import re

CUSTOMERS = {'buzzbee', 'dickie', 'caixing', 'edu', '360', 'yinhui', 'disney', 'maxx', 'shushupapa', 'barter'}
TARGET_TEMPLATE = 'HEYUAN_BUSINESS_UNIFIED_HUAXING_V2'
BUZZBEE_SHEETS = ('ITEM表-子弹枪', 'ITEM表-水枪', 'ITEM表-套装', 'ITEM表-公仔')
DICKIE_EXTRA = '恐龙蛋口水车Iteam表'
ADVISORY_CUSTOMERS = {'dickie', 'seasons', 'maxx', 'shushupapa', 'barter'}
DUPLICATE_CODES = {'duplicate_reference', 'existing_order_line', 'duplicate_batch_order_line', 'duplicate_existing_order'}


def apply_advisory_policy(rows):
    """Business-approved Huaxing warnings; identity and quantity conflicts stay hard."""
    from app.services.customer_order_unified import _refresh_status, _text, _number, _issue
    for row in rows:
        if _number(row.get('quantity')) is None or _number(row.get('quantity')) <= 0:
            if not any(i.get('field') == 'quantity' and i.get('severity') == 'blocked' for i in row.get('issues', [])):
                row.setdefault('issues', []).append(_issue('blocked', 'missing_quantity', 'quantity', '缺少有效订单数量，请核对原单'))
        for field, label in (('date_code', '日期码'), ('units_per_carton', '装箱数')):
            if not _text(row.get(field)) and not any(i.get('field') == field for i in row.get('issues', [])):
                issue = _issue('warning', 'missing_' + field, field, f'{label}暂缺，可补录；留空仍可导出')
                issue.update(can_edit=True, can_skip=True, edit_field=field, edit_label=label,
                             edit_input_type='text' if field == 'date_code' else 'number')
                row.setdefault('issues', []).append(issue)
        for issue in row.get('issues', []):
            if issue.get('code') in DUPLICATE_CODES:
                issue.update(code='duplicate_order_warning', severity='warning', can_skip=False,
                             skip_key='', skip_label='')
                issue['message'] = issue['message'].rstrip('；。') + '；导出新行以红色字体标记，请核对重复订单'
            elif issue.get('code') in {'past_ship_date', 'past_inspection_date', 'missing_po_unit_prices'}:
                from app.services.customer_order_manual import MANUAL_EDITABLE_FIELDS
                issue.update(severity='warning', can_skip=True)
                field = issue.get('edit_field') or issue.get('field')
                if field in MANUAL_EDITABLE_FIELDS:
                    meta = MANUAL_EDITABLE_FIELDS[field]
                    issue.update(can_edit=True, edit_field=field, edit_label=meta['label'], edit_input_type=meta['input_type'])
            elif issue.get('severity') == 'blocked':
                issue.update(requires_source_correction=True, can_skip=False, can_edit=False)
        _refresh_status(row)


def _extension_column(sheet, header, *, after=19):
    from copy import copy
    from app.services.customer_order_unified import _normalized_header
    matches = [c.column for c in sheet[3] if _normalized_header(c.value) == header]
    if len(matches) > 1:
        from app.services.customer_order_unified import CustomerOrderUnifiedError
        raise CustomerOrderUnifiedError(f'{sheet.title} 的“{header}”列重复，请先核对表头')
    if matches:
        return matches[0]
    column = max(after, sheet.max_column) + 1
    target, source = sheet.cell(3, column), sheet.cell(3, after)
    target._style = copy(source._style)
    target.value = header
    from openpyxl.utils import get_column_letter
    sheet.column_dimensions[get_column_letter(column)].width = 30 if header == 'PO单价明细' else 22
    return column


def _mark_generated_row(sheet, number, *, duplicate=False, blue=False):
    if not duplicate and not blue:
        return
    from copy import copy
    from openpyxl.styles import Font
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.utils import get_column_letter
    color = 'FFFF0000' if duplicate else 'FF0000FF'
    for cell in sheet[number]:
        font = copy(cell.font)
        font.color = color
        cell.font = font
    # Only the generated row bypasses legacy duplicate rules (some use wrong
    # columns or #REF!). Historical conditional formatting is left intact.
    for rules in sheet.conditional_formatting._cf_rules.values():
        for rule in rules:
            rule.priority += 1
    rule = FormulaRule(formula=['TRUE'], font=Font(color=color), stopIfTrue=True)
    rule.priority = 1
    sheet.conditional_formatting.add(f'A{number}:{get_column_letter(sheet.max_column)}{number}', rule)


def _repair_dickie_duplicate_rules(workbook):
    """Correct only the supplied SO/item duplicate rules, preserving other styles."""
    def legacy_pattern(ref, item):
        return re.compile(
            rf'AND\((?:\${ref}\d+|#REF!)<>"",(?:\${item}\d+|#REF!)<>"",'
            rf'COUNTIFS\(\${ref}\$\d+:\${ref}\$\d+,(?:\${ref}\d+|#REF!),'
            rf'\${item}\$\d+:\${item}\$\d+,(?:\${item}\d+|#REF!)\)>1\)')
    def guarded(ref, item, row, end):
        return (f'AND(LEFT(UPPER(TRIM(${ref}{row})),2)="SC",${item}{row}<>"",'
                f'ISERROR(SEARCH("MO",$B{row})),ISERROR(SEARCH("备料",$B{row})),'
                f'COUNTIFS(${ref}$4:${ref}${end},${ref}{row},'
                f'${item}$4:${item}${end},${item}{row},'
                f'$B$4:$B${end},"<>*MO*",$B$4:$B${end},"<>*备料*")>1)')
    # Recognize our own exact formula shape on re-export, not arbitrary rules.
    for name in ('接单表', '正单评审表', 'ITEM表', DICKIE_EXTRA):
        if name not in workbook.sheetnames:
            continue
        sheet = workbook[name]
        ref_col, item_col = ('D', 'G') if name in ('接单表', '正单评审表') else ('E', 'H')
        patterns = [legacy_pattern('E', 'H')]
        if name in ('接单表', '正单评审表'):
            patterns.append(legacy_pattern('D', 'G'))
        def row_shape(formula):
            return re.sub(r'(\$?[A-Z]{1,3}\$?)\d+', r'\1#', formula)
        shape = row_shape(guarded(ref_col, item_col, 4, 100))
        for key, rules in sheet.conditional_formatting._cf_rules.items():
            start = min(cell_range.min_row for cell_range in key.sqref.ranges)
            end = max(sheet.max_row, start)
            for rule in rules:
                if rule.type != 'expression' or len(rule.formula or []) != 1:
                    continue
                formula = rule.formula[0]
                own_shape = row_shape(formula)
                if any(pattern.fullmatch(formula) for pattern in patterns) or own_shape == shape:
                    rule.formula = [guarded(ref_col, item_col, start, end)]


def enabled(factory_id, customer_code):
    return (factory_id == 'huaxing' and customer_code in CUSTOMERS) or (
        factory_id == 'huakang-d' and customer_code == 'disney'
    )


def target_template(factory_id):
    return 'HEYUAN_BUSINESS_UNIFIED_HUAKANG_D_DISNEY_V1' if factory_id == 'huakang-d' else TARGET_TEMPLATE


def item_sheets(workbook, customer_code):
    if customer_code == 'buzzbee' and any(name in workbook.sheetnames for name in BUZZBEE_SHEETS):
        return list(BUZZBEE_SHEETS)
    return ['ITEM表'] + ([DICKIE_EXTRA] if customer_code == 'dickie' and DICKIE_EXTRA in workbook.sheetnames else [])


def outer_pack(value):
    from app.services.customer_order_unified import _text, _number, _number_text
    text = _text(value)
    if re.fullmatch(r'\d+\s*/\s*\d+', text):
        text = text.rsplit('/', 1)[1]
    number = _number(text)
    return _number_text(number) if number is not None and number > 0 else ''


def extra_columns(sheet):
    from app.services.customer_order_unified import _normalized_header
    result = {}
    for field, header in (('remarks', '备注'), ('contact', '联系人'), ('pdq', 'PDQ')):
        columns = [cell.column for cell in sheet[3] if cell.column > 28 and _normalized_header(cell.value) == header]
        if len(columns) == 1:
            result[field] = columns[0]
    return result


def edu_order_key(contract, po):
    from app.services.customer_order_unified import _key, _text
    return _key(contract), _key(re.sub(r'^PO\s*#?\s*', '', _text(po), flags=re.I))


def edu_identity(contract, po, product):
    from app.services.customer_order_unified import _key
    from app.services.huaxing_order_legacy.edu_schedule import normalize_edu_item
    order = edu_order_key(contract, po)
    return (*order, _key(normalize_edu_item(product))) if any(order) else ('', '', '')


def edu_history_matches(row, history):
    contract, po = edu_order_key(row.get('contract_no'), row.get('po_no'))
    if not contract and not po:
        return []
    # A complete composite contract must match completely. Legacy single-part
    # contracts may reconcile to a unique composite, never to a sibling suffix.
    return [h for h in history if edu_order_key(h.contract_no, h.po_no)[1] == po and
            (edu_order_key(h.contract_no, h.po_no)[0] == contract or
             contract and '/' not in contract and contract in
             {edu_order_key(part, '')[0] for part in h.contract_no.split('/')})]


def allocate_edu_references(records, next_number):
    assigned = {}
    for record in records:
        order = edu_identity(record.get('contract_no'), record.get('customer_po'), record.get('item_no'))
        if record.get('huaxing_po'):
            continue
        if order not in assigned:
            assigned[order] = f'EDUHX{next_number:05d}'
            next_number += 1
        record['huaxing_po'] = assigned[order]


def read_history(workbook, customer_code):
    from app.services.customer_order_unified import UnifiedHistoryRow, _text
    from app.services.customer_order_regional import history_number
    fields = {1: 'received_date', 2: 'order_type', 3: 'production_no', 4: 'contract_no',
              5: 'reference_no', 6: 'po_no', 7: 'customer_name', 8: 'product_no',
              9: 'product_name_zh', 10: 'product_name_en', 14: 'standard', 20: 'packaging',
              21: 'date_code', 26: 'requested_ship_date'}
    result = []
    for name in item_sheets(workbook, customer_code):
        sheet = workbook[name]
        extras = extra_columns(sheet)
        for r in range(4, sheet.max_row + 1):
            values = {field: _text(sheet.cell(r, col).value) if sheet.cell(r, col).data_type != 'f' else '' for col, field in fields.items()}
            remarks = _text(sheet.cell(r, extras['remarks']).value) if 'remarks' in extras else ''
            reference = values['reference_no'] or values['contract_no'] or values['po_no']
            if not values['product_no'] or values['product_no'] == '产品编号' or '模拟数据' in remarks:
                continue
            values['reference_no'] = reference
            if customer_code == 'edu':
                from app.services.huaxing_order_legacy.edu_schedule import normalize_edu_item
                from app.services.huaxing_order_legacy.edu_grouped_export import is_reservation_row
                values['product_no'] = normalize_edu_item(values['product_no'])
                if is_reservation_row(sheet, r):
                    # Reservations enrich product fields, but are not customer orders.
                    values.update(reference_no='', contract_no='', po_no='')
            values['quantity'] = history_number(sheet.cell(r, 11), workbook, customer_code)
            values['units_per_carton'] = outer_pack(history_number(sheet.cell(r, 12), workbook, customer_code, pack=True))
            result.append(UnifiedHistoryRow(row=r, source_sheet=name, **values,
                extras={field: _text(sheet.cell(r, col).value) for field, col in extras.items() if sheet.cell(r, col).data_type != 'f'}))
    return result


def map_record(row, record, customer_code):
    from app.services.customer_order_unified import _text
    get = lambda field: _text(record.get(field))
    if customer_code == 'edu':
        row.update(reference_no=get('huaxing_po'), production_no='',
                   contract_no=get('contract_no'), po_no=get('customer_po'),
                   source_product_no=get('source_item_no') or get('item_no'),
                   source_order_date=get('order_date'))
    elif customer_code == '360':
        row.update(reference_no=get('production_no'), contract_no='',
                   po_no=get('unified_customer_po') or get('customer_po') or get('po_no'), product_no=get('item_full') or row.get('product_no', ''),
                   packaging=get('packaging'), standard=get('standard'))
    elif customer_code == 'yinhui':
        so, po = get('so_no'), get('contract_no')
        row.update(reference_no=re.sub(r'^SO\s*', '', so, flags=re.I), contract_no='',
                   po_no=re.sub(r'^PO\s*', '', po, flags=re.I), packaging=get('packaging'), standard=get('standard'))
    elif customer_code == 'disney':
        row.update(reference_no=get('po_number'), contract_no='', standard=get('standard'))
    elif customer_code in {'maxx', 'shushupapa', 'barter'}:
        row.update(reference_no=get('po_number'), packaging=get('color_box'), standard=get('standard'),
                   units_per_carton=outer_pack(get('outer_pack')))
        if customer_code in {'shushupapa', 'barter'}:
            row['po_no'] = get('customer_po')
    if customer_code in {'360', 'yinhui'} and get('product_name') and not re.search(r'[\u3400-\u9fff]', get('product_name')):
        row['product_name_en'] = get('product_name')
        row['product_name_zh'] = ''
    for field in ('remarks', 'contact', 'pdq'):
        if get(field):
            row[field] = get(field)
    if get('battery'):
        row['remarks'] = '\n'.join(filter(None, [row.get('remarks'), '电池：' + get('battery')]))
    if get('box_mark'):
        row['carton_mark'] = get('box_mark')
    return row


def prepare_row(row, history, customer_code, available_sheets, factory_id='huaxing'):
    from app.services.customer_order_unified import _key, _text, _issue
    row['_huaxing_customer'] = customer_code
    if customer_code == 'edu':
        from app.services.huaxing_order_legacy.edu_schedule import normalize_edu_item
        row['product_no'] = normalize_edu_item(row.get('product_no'))
    if customer_code == 'caixing':
        row['reference_no'], row['contract_no'] = row.get('contract_no', ''), ''
    if customer_code == 'dickie':
        row['production_no'], row['contract_no'] = row.get('contract_no', ''), ''
    matches = [h for h in history if _key(h.product_no) == _key(row.get('product_no'))]
    if customer_code in {'barter', 'maxx', 'shushupapa'} and not _text(row.get('date_code')):
        same_order = [h for h in matches if h.reference_no and _key(h.reference_no) == _key(row.get('reference_no'))]
        codes = {h.date_code for h in same_order if h.date_code}
        if len(codes) == 1:
            row['date_code'] = next(iter(codes))
            row.setdefault('lineage', {})['date_code'] = '当前排期同合同/同产品唯一日期码；原单未可靠识别'
            row['issues'] = [i for i in row.get('issues', []) if i.get('code') != 'missing_date_code']
    if customer_code == 'buzzbee' and not matches:
        # Latest BuzzBee item suffixes describe schedule variants, not PO line numbers.
        base = _key(row.get('product_no'))
        candidates = [h for h in history if _key(h.product_no.split('-', 1)[0]) == base]
        same_order = [h for h in candidates if h.reference_no and _key(h.reference_no) == _key(row.get('contract_no'))]
        same_customer = [h for h in candidates if _key(h.customer_name) and _key(h.customer_name) == _key(row.get('customer_name'))]
        evidence = same_order or same_customer
        if not evidence and candidates:
            row.setdefault('issues', []).append(_issue('blocked', 'ambiguous_product_variant', 'product_no',
                '基础货号带排期版本后缀，但没有同订单或同客户依据，请核对完整货号'))
        candidates = evidence or candidates
        products = {h.product_no for h in candidates}
        if len(products) == 1 and evidence:
            original = row['product_no']
            row['product_no'] = next(iter(products))
            row.setdefault('lineage', {})['product_no'] = f'PO {original} → 最新排期唯一订单/客户货号 {row["product_no"]}'
            matches = candidates
        elif candidates:
            matches = candidates
        if len(products) > 1 and evidence:
            row.setdefault('issues', []).append(_issue('blocked', 'ambiguous_product_variant', 'product_no',
                '相同基础货号在排期有多个版本，无法根据订单或客户唯一确定'))
    if customer_code == '360':
        full_item = _text(row.get('product_no'))
        if re.fullmatch(r'\d+\.\d+', full_item):
            base_item = full_item.split('.', 1)[0]
            base_matches = [h for h in history if _key(h.product_no) == _key(base_item)]
            if base_matches:
                row['product_no'] = base_item
                matches = base_matches
        reference = _text(row.get('reference_no'))
        same_release = [h for h in matches if re.fullmatch(re.escape(reference) + r'(?:-\d+)?', h.reference_no)]
        po_values = {h.po_no for h in same_release if _key(row.get('po_no')) and _key(h.po_no).endswith(_key(row['po_no']))}
        if len(po_values) == 1:
            row['po_no'] = next(iter(po_values))
        refs = {h.reference_no for h in same_release}
        if len(refs) == 1:
            row['reference_no'] = next(iter(refs))
        elif len(refs) > 1:
            row.setdefault('issues', []).append(_issue('blocked', 'existing_split_release', 'row',
                '该 Release 在最新排期已拆成多个子单，请核对拆单数量后录入'))
    if customer_code == 'edu':
        if not any(edu_order_key(row.get('contract_no'), row.get('po_no'))):
            row.setdefault('issues', []).append(_issue('blocked', 'edu_missing_business_identity', 'contract_no',
                'EDU 缺少客户合同号或客户 PO；自动分配的 EDUHX 编号不能替代客户订单身份'))
        same_order = edu_history_matches(row, matches)
        refs = {h.reference_no for h in same_order}
        contracts = {h.contract_no for h in same_order}
        if len(refs) == 1 and len(contracts) == 1:
            row['reference_no'] = next(iter(refs))
            row['contract_no'] = next(iter(contracts))
            packs = {outer_pack(h.units_per_carton) for h in same_order if outer_pack(h.units_per_carton)}
            if not outer_pack(row.get('units_per_carton')) and len(packs) == 1:
                row['units_per_carton'] = next(iter(packs))
                row.setdefault('lineage', {})['units_per_carton'] = '当前排期同客户合同/PO与货号的唯一装箱数'
        elif same_order:
            row.setdefault('issues', []).append(_issue('blocked', 'ambiguous_edu_history', 'contract_no',
                '客户合同/PO 与产品对应多个历史编号，必须核对完整订单身份'))
    lanes = {h.source_sheet for h in matches}
    if len(available_sheets) == 1:
        lane = available_sheets[0]
    elif len(lanes) == 1:
        lane = next(iter(lanes))
    else:
        lane = ''
        row.setdefault('issues', []).append(_issue('blocked', 'unmapped_item_sheet', 'item_sheet_name',
            '货号无法唯一匹配最新排期的 ITEM 分类页，请核对分类后再录入'))
    row['item_sheet_name'] = lane
    row['units_per_carton'] = outer_pack(row.get('units_per_carton'))
    factory_name = '华康D' if factory_id == 'huakang-d' else '华兴'
    row.setdefault('lineage', {})['mapping_rule'] = f'{factory_name}最新统一排期：按实际 ITEM 分类页及公共字段写入，摘要关联对应明细行。'
    for field in ('contact', 'pdq'):
        values = {h.extras[field] for h in matches if h.source_sheet == lane and h.extras.get(field)}
        if not _text(row.get(field)) and len(values) == 1:
            row[field] = next(iter(values))


def write_extras(sheet, row_number, row):
    for field, column in extra_columns(sheet).items():
        sheet.cell(row_number, column).value = row.get(field) or None


def validate_edited_identities(preview, original, history, available_sheets, overrides=()):
    from app.services import customer_order_unified as unified
    for row in preview['rows']:
        current = tuple(unified._text(row.get(k)) for k in ('product_no', 'reference_no', 'contract_no', 'po_no'))
        if current == original[row['id']]:
            continue
        explicit = {o['field'] for o in overrides if o['row_id'] == row['id']}
        if current[0] != original[row['id']][0]:
            for field in ('product_name_zh', 'product_name_en', 'units_per_carton', 'standard', 'packaging', 'unit_price_hkd'):
                lineage = str(row.get('lineage', {}).get(field, ''))
                if field not in explicit and ('历史' in lineage or '当前' in lineage and '排期' in lineage):
                    row[field] = unified._unique_history_value(history, row.get('product_no'), field, row.get('customer_name', ''))
                    row['lineage'][field] = '人工更正货号后重新核对最新排期唯一历史值'
        matches = [h for h in history if unified._key(h.product_no) == unified._key(row.get('product_no'))]
        lanes = {h.source_sheet for h in matches}
        if len(available_sheets) == 1:
            row['item_sheet_name'] = available_sheets[0]
        elif len(lanes) == 1:
            row['item_sheet_name'] = next(iter(lanes))
        else:
            raise unified.CustomerOrderUnifiedError('修改后的货号无法唯一确定 ITEM 分类页，请修订 PO 后重新预览')
        same_order = [h for h in matches if h.reference_no and unified._key(h.reference_no) == unified._key(row.get('reference_no'))]
        batch_matches = [other for other in preview['rows'] if other is not row and
                         unified._key(other.get('product_no')) == unified._key(row.get('product_no')) and
                         unified._key(other.get('reference_no')) == unified._key(row.get('reference_no'))]
        if row.get('_huaxing_customer') == 'edu':
            if not any(edu_order_key(row.get('contract_no'), row.get('po_no'))):
                raise unified.CustomerOrderUnifiedError('EDU 修改后缺少客户合同号或客户 PO，请重新核对')
            same_order += edu_history_matches(row, matches)
            identity = edu_identity(row.get('contract_no'), row.get('po_no'), row.get('product_no'))
            batch_matches += [other for other in preview['rows'] if other is not row and
                              edu_identity(other.get('contract_no'), other.get('po_no'), other.get('product_no')) == identity]
        if same_order or batch_matches:
            raise unified.CustomerOrderUnifiedError('修改后的订单身份存在重复或数量冲突，原重复确认不适用，请修订 PO 后重新预览')


def _repair_shushupapa_summary_amounts(workbook, order_row, review_row, row):
    """Bind new detail amounts to their prices, never a copied subtotal row."""
    from app.services import customer_order_unified as unified

    order = workbook['接单表']
    review = workbook['正单评审表']
    for column, field in ((14, 'unit_price_hkd'), (16, 'unit_price_usd')):
        order.cell(order_row, column).value = unified._number(row.get(field))
        letter = 'N' if column == 14 else 'P'
        ref = f"'接单表'!{letter}{order_row}"
        review.cell(review_row, column).value = f'=IF(LEN({ref})=0,"",{ref})'

    for sheet, number in ((order, order_row), (review, review_row)):
        for amount_column, price_column in ((15, 'N'), (17, 'P')):
            sheet.cell(number, amount_column).value = (
                f'=IF(OR(LEN($J{number})=0,LEN(${price_column}{number})=0),'
                f'"",ROUND($J{number}*${price_column}{number},2))'
            )
        for column in range(14, 18):
            sheet.cell(number, column).number_format = '#,##0.00'


def export_rows(workbook, rows):
    if rows and all(row.get('_huaxing_customer') == 'edu' for row in rows):
        from app.services.huaxing_order_legacy.edu_grouped_export import export_rows as export_edu
        return export_edu(workbook, rows)
    from app.services import customer_order_unified as unified
    from app.services.huakang_a_order_legacy.unified_schedule import output_slots
    for row in rows:
        if row.get('item_sheet_name') not in item_sheets(workbook, row['_huaxing_customer']):
            raise unified.CustomerOrderUnifiedError('尚未确定产品的 ITEM 分类页，请核对货号及分类后重新预览')
    buzzbee_column = None
    if any(row.get('_huaxing_customer') == 'buzzbee' for row in rows):
        from app.services.customer_order_buzzbee_color_box import prepare_review
        buzzbee_column = prepare_review(workbook)
    dickie = bool(rows) and all(row.get('_huaxing_customer') == 'dickie' for row in rows)
    grouped_handles = None
    if dickie:
        from app.services.customer_order_dickie_layout import allocate_rows
        grouped_handles = allocate_rows(workbook, rows)
        _repair_dickie_duplicate_rules(workbook)
    else:
        counts = Counter(row['item_sheet_name'] for row in rows)
        counts.update({'接单表': len(rows), '正单评审表': len(rows)})
        slots = output_slots(
            workbook, len(rows), sheet_counts=counts,
            top_order_area=any(row.get('_huaxing_customer') == 'shushupapa' for row in rows),
        )
    used = Counter()
    dickie_columns = {}
    if any(row.get('_huaxing_customer') == 'dickie' for row in rows):
        for name in ('接单表', '正单评审表'):
            dickie_columns[name] = {header: _extension_column(workbook[name], header)
                for header in ('生产单号', 'PO单价明细', '单价口径')}
    for index, row in enumerate(rows):
        name = row['item_sheet_name']
        item_row = grouped_handles[index][name].row if grouped_handles is not None else slots[name][used[name]]
        used[name] += 1
        unified._write_item_row(workbook[name], item_row, row)
        write_extras(workbook[name], item_row, row)
        for summary in ('接单表', '正单评审表'):
            summary_row = grouped_handles[index][summary].row if grouped_handles is not None else slots[summary][index]
            unified._write_summary_row(workbook[summary], summary_row, row,
                                       item_row=item_row, item_sheet=workbook[name])
            if row.get('_huaxing_customer') == 'dickie':
                sheet, number = workbook[summary], summary_row
                cols = dickie_columns[summary]
                ref = f"'{name}'!C{item_row}"
                sheet.cell(number, cols['生产单号']).value = f'=IF(LEN({ref})=0,"",{ref})'
                sheet.cell(number, cols['PO单价明细']).value = row.get('price_details_text') or None
                sheet.cell(number, cols['单价口径']).value = '按PO数量加权均价' if row.get('price_basis') == 'weighted' else '原单单价'
                if row.get('price_details'):
                    from copy import copy
                    cell = sheet.cell(number, cols['PO单价明细'])
                    alignment = copy(cell.alignment)
                    alignment.wrapText = True
                    cell.alignment = alignment
                    sheet.cell(number, 14).value = unified._number(row.get('unit_price_hkd'))
                    sheet.cell(number, 14).number_format = '0.00#####'
                    sheet.cell(number, 15).value = unified._number(row.get('amount_hkd'))
                    sheet.cell(number, 15).number_format = '0.00'
                    unified._fit_generated_text(sheet, number)
        if row.get('_huaxing_customer') == 'buzzbee':
            from app.services.customer_order_buzzbee_color_box import write_new_row
            write_new_row(workbook, buzzbee_column, slots['正单评审表'][index], workbook[name], item_row)
        if row.get('_huaxing_customer') == 'shushupapa':
            _repair_shushupapa_summary_amounts(
                workbook, slots['接单表'][index], slots['正单评审表'][index], row,
            )
        duplicate = any(i.get('code') in DUPLICATE_CODES | {'duplicate_order_warning'} for i in row.get('issues', []))
        blue = row.get('_huaxing_customer') == 'dickie'
        _mark_generated_row(workbook[name], item_row, duplicate=duplicate, blue=blue)
        for summary in ('接单表', '正单评审表'):
            summary_row = grouped_handles[index][summary].row if grouped_handles is not None else slots[summary][index]
            _mark_generated_row(workbook[summary], summary_row, duplicate=duplicate, blue=blue)
        # Some supplied EDU reserved rows have a broken price lookup anchor.
        # New paired summary rows can reference their exact order counterpart;
        # the order sheet's manually maintained price remains authoritative.
        review_row = grouped_handles[index]['正单评审表'].row if grouped_handles is not None else slots['正单评审表'][index]
        order_row = grouped_handles[index]['接单表'].row if grouped_handles is not None else slots['接单表'][index]
        review_price = workbook['正单评审表'].cell(review_row, 14)
        if isinstance(review_price.value, str) and '#REF!' in review_price.value and '接单表!' in review_price.value:
            ref = f"'接单表'!N{order_row}"
            review_price.value = f'=IF(LEN({ref})=0,"",{ref})'
