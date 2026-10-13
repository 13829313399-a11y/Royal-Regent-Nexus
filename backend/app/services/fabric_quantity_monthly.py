"""Internal quantity reconciliation across all actual fabric receiving/issuing channels.

This is a read of immutable stock evidence, never a second stock posting or a
supplier portal ownership decision. Quantities use their own business dates.
"""
import json
from collections import defaultdict
from decimal import Decimal
from io import BytesIO

from fastapi import HTTPException
from sqlalchemy import select

from app.models.fabric_master import FabricMasterRecord
from app.models.fabric_receiving import FabricReceipt
from app.services import warehouse_operations as operations
from app.services.material_handover_monthly import bounds
from app.services.fabric_procurement_parser import digest

FIELDS = ('received', 'receipt_reversed', 'net_received', 'issued', 'issue_reversed',
          'returned', 'net_issued', 'net_outbound', 'stock_change',
          'manual_received', 'scan_received', 'manual_outbound', 'scan_outbound')
CHANNELS = {'MANUAL': '手动', 'SCAN': '扫码'}


def report(db, month, direction='', channel='', party_key=''):
    operations.ensure_schema(db)
    start, end = bounds(month)
    if direction not in ('', 'IN', 'OUT') or channel not in ('', *CHANNELS):
        raise HTTPException(422, '收发方向或登记方式不正确')
    operations.procurement.lock_factory(db)
    _, projected = operations.projection(db, 'fabric')
    entries = list(projected.values())
    # Only unambiguous, exact recorded names join internal supplier groups.
    # This does not bind old documents or grant an external token ownership.
    aliases = defaultdict(set)
    names = {}
    for master in db.scalars(select(FabricMasterRecord).where(
            FabricMasterRecord.factory_id == operations.FACTORY, FabricMasterRecord.kind == 'SUPPLIER')):
        names[master.id] = master.name
        aliases[master.name].add(master.id)
    for entry in entries:
        body = entry['data']
        if entry['kind'] == 'RECEIPT' and body.get('handover_supplier_id'):
            aliases[body['counterparty']].add(body['handover_supplier_id'])
            names.setdefault(body['handover_supplier_id'], body['counterparty'])

    events = []

    def event(identifier, when, body, kind, sign=1, original_id='', actor=''):
        if not start <= when <= end:
            return
        supplier = kind == 'RECEIPT'
        recorded_party = body.get('counterparty', '')
        supplier_id = body.get('handover_supplier_id', '') if supplier else ''
        matched = aliases.get(recorded_party, set())
        if supplier and not supplier_id and len(matched) == 1:
            supplier_id = next(iter(matched))
            basis = '单据名称唯一匹配（仅内部汇总）'
        else:
            basis = '供应商基础资料编号' if supplier_id else '原单往来名称'
        identity = ('supplier:' + supplier_id if supplier_id else 'supplier-name:' + recorded_party) if supplier else 'recipient:' + recorded_party
        method = 'SCAN' if body.get('handover_document_id') else 'MANUAL'
        prices = body.get('handover_line_snapshot', {})
        events.append(dict(id=identifier, date=when, direction='IN' if supplier else 'OUT', kind=kind,
            correction=sign < 0, original_id=original_id, channel=method, channel_label=CHANNELS[method],
            origin=body.get('source_system') or 'MANUAL', source_no=body.get('source_no') or original_id or identifier,
            source_line=body.get('source_line', ''), document_id=body.get('handover_document_id', ''),
            party_key=identity, party_name=names.get(supplier_id, recorded_party) if supplier else recorded_party,
            recorded_party=recorded_party, supplier_id=supplier_id, identity_basis=basis,
            item_code=body.get('item_code', ''), item_name=body.get('item_name', ''), unit=body.get('unit', ''),
            quantity=str(Decimal(body['quantity']) * sign), actor_name=actor,
            unit_price=prices.get('unit_price', body.get('unit_price', '')),
            price_unit=prices.get('price_unit', body.get('price_unit', '')),
            currency=prices.get('currency', body.get('currency', '')),
            price_note=prices.get('price_note', body.get('price_note', ''))))
        events[-1]['amount'] = body.get('amount', '')

    # One physical receipt header once, irrespective of its roll/bin split.
    for receipt in db.scalars(select(FabricReceipt).where(FabricReceipt.factory_id == operations.FACTORY,
            FabricReceipt.receipt_date >= start, FabricReceipt.receipt_date <= end)):
        facts = json.loads(receipt.source_json)
        pricing = json.loads(receipt.payload_json)
        event(receipt.id, receipt.receipt_date, dict(counterparty=facts.get('supplier', ''),
            item_code=facts.get('material_code', ''), item_name=facts.get('material_name', ''),
            quantity=receipt.quantity, unit=receipt.unit, source_no=pricing.get('delivery_reference', receipt.delivery_reference),
            source_line=receipt.source_line_id, source_system='PROCUREMENT_RECEIVING',
            unit_price=pricing.get('unit_price', facts.get('unit_price') or ''), price_unit=pricing.get('price_unit', ''),
            currency=pricing.get('currency', ''), amount=pricing.get('amount', ''),
            price_note=pricing.get('price_basis', '')), 'RECEIPT', actor=receipt.actor_name)
    for entry in entries:
        kind, body = entry['kind'], entry['data']
        if kind in ('RECEIPT', 'ISSUE', 'RETURN'):
            # A workshop return follows its original issue's party/channel,
            # and is shown separately; it never reduces supplier receipts.
            if kind == 'RETURN':
                original = projected[body['original_id']]['data']
                body = {**body, **{key: original.get(key, '') for key in
                    ('handover_document_id', 'handover_supplier_id')}}
            event(entry['id'], entry['business_date'], body, kind,
                  original_id=body.get('original_id', ''), actor=entry['actor_name'])
        elif kind == 'CORRECTION':
            if body.get('procurement_original_receipt_id'):
                event(entry['id'], entry['business_date'], body, 'RECEIPT', -1,
                    original_id=body['procurement_original_receipt_id'], actor=entry['actor_name'])
                continue
            original = projected.get(body['original_id'])
            if original and original['kind'] in ('RECEIPT', 'ISSUE'):
                event(entry['id'], entry['business_date'], original['data'], original['kind'], -1,
                      original_id=original['id'], actor=entry['actor_name'])
        # Requests, transfers, metadata and unused printed documents are not actual handovers.

    parties = {row['party_key']: dict(key=row['party_key'], name=row['party_name'], direction=row['direction'])
               for row in events if not direction or row['direction'] == direction}
    details = [row for row in events if (not direction or row['direction'] == direction)
               and (not channel or row['channel'] == channel) and (not party_key or row['party_key'] == party_key)]
    groups, units = {}, {}

    def add(total, row):
        quantity = Decimal(row['quantity'])
        if row['kind'] == 'RECEIPT':
            total['receipt_reversed' if row['correction'] else 'received'] += -quantity if row['correction'] else quantity
            total['net_received'] += quantity
            total[row['channel'].lower() + '_received'] += quantity
            total['stock_change'] += quantity
        elif row['kind'] == 'ISSUE':
            total['issue_reversed' if row['correction'] else 'issued'] += -quantity if row['correction'] else quantity
            total['net_issued'] += quantity
            total['net_outbound'] += quantity
            total[row['channel'].lower() + '_outbound'] += quantity
            total['stock_change'] -= quantity
        else:
            total['returned'] += quantity
            total['net_outbound'] -= quantity
            total[row['channel'].lower() + '_outbound'] -= quantity
            total['stock_change'] += quantity

    for row in details:
        key = (row['direction'], row['party_key'], row['item_code'], row['unit'])
        group = groups.setdefault(key, {**{k: row[k] for k in
            ('direction', 'party_key', 'party_name', 'item_code', 'item_name', 'unit')},
            **dict.fromkeys(FIELDS, Decimal(0)), 'document_ids': set()})
        add(group, row)
        group['document_ids'].add(row['original_id'] if row['correction'] else row['id'])
        total = units.setdefault(row['unit'], {'unit': row['unit'], **dict.fromkeys(FIELDS, Decimal(0))})
        add(total, row)
    summary = [{**{k: str(v) if isinstance(v, Decimal) else v for k, v in group.items() if k != 'document_ids'},
                'document_count': len(group['document_ids'])} for group in groups.values()]
    totals = [{k: str(v) if isinstance(v, Decimal) else v for k, v in group.items()} for group in units.values()]
    result = dict(month=month, direction=direction, channel=channel, party_key=party_key,
        parties=sorted(parties.values(), key=lambda r: (r['direction'], r['name'], r['key'])),
        details=sorted(details, key=lambda r: (r['date'], r['id'])),
        summary=sorted(summary, key=lambda r: (r['direction'], r['party_name'], r['item_code'], r['unit'])),
        totals=sorted(totals, key=lambda r: r['unit']), cutting_interface='PENDING', live=True,
        closed=False, amounts_calculated=False,
        basis='按实收、实发、车间退料及更正的业务日期；手动和扫码合并一次，按往来方、物料和单位分别核对')
    result['version'] = digest(result)
    return result


def export_xlsx(result):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    pages = [
        ('收发数量汇总', result['summary'], [('direction', '方向'), ('party_name', '供应商 / 领用方'),
            ('item_code', '物料编码'), ('item_name', '名称'), ('unit', '单位'), ('document_count', '凭证数')]),
        ('收发原单明细', result['details'], [('date', '业务日期'), ('direction', '方向'), ('party_name', '归属方'),
            ('recorded_party', '原单往来方'), ('identity_basis', '分组依据'), ('channel_label', '登记方式'),
            ('origin', '来源'), ('source_no', '原单号'), ('source_line', '明细号'), ('id', '库存凭证'),
            ('original_id', '关联原凭证'), ('kind', '业务'), ('correction', '是否冲销'),
            ('item_code', '物料编码'), ('item_name', '名称'), ('quantity', '数量（冲销为负）'), ('unit', '单位'),
            ('unit_price', '参考单价'), ('price_unit', '计价单位'), ('currency', '币种'), ('amount', '原入库单金额'), ('price_note', '价格依据')]),
    ]
    labels = dict(received='实收', receipt_reversed='收料冲销', net_received='净实收', issued='实发',
        issue_reversed='发料冲销', returned='车间退料', net_issued='实发减冲销', net_outbound='净发料（扣车间退料）',
        stock_change='本月库存增减', manual_received='手动净实收', scan_received='扫码净实收',
        manual_outbound='手动净发料', scan_outbound='扫码净发料')
    for title, rows, columns in pages:
        if title == '收发数量汇总':
            columns += [(field, labels[field]) for field in FIELDS]
        sheet = book.create_sheet(title)
        sheet.append([result['month'], '实时数量核对；未关账；参考价格不计算应付金额', result['version']])
        sheet.append([label for _, label in columns])
        for row in rows:
            sheet.append([str(row.get(key, '')) for key, _ in columns])
        sheet.freeze_panes = 'A3'
        sheet.auto_filter.ref = f'A2:{get_column_letter(sheet.max_column)}{sheet.max_row}'
        for cells in sheet:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for cell in sheet[2]:
            cell.font = Font(bold=True)
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 22
    output = BytesIO()
    book.save(output)
    return output.getvalue()
