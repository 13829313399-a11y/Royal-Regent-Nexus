from __future__ import annotations

from copy import copy
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
import math
import unicodedata
from tempfile import TemporaryDirectory
from typing import Any, Callable

import openpyxl
from app.services.sparse_worksheet import insert_rows as insert_sparse_rows
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula
from app.services.huakang_a_order_legacy import unified_schedule as huakang_unified
from app.services import customer_order_regional as regional
from app.services import customer_order_huaxing_unified as huaxing

from app.services.customer_order_buzzbee import (
    CustomerOrderWorkbookError,
    MAX_BATCH_PO_BYTES,
    MAX_BATCH_PO_FILES,
    MAX_PO_BYTES,
    MAX_SCHEDULE_BYTES,
)
from app.services.customer_order_manual import (
    apply_overrides_to_preview,
    decorate_manual_resolution_policy,
)


PREVIEW_SCHEMA_VERSION = "customer-order-heyuan-unified-preview-v1"
TARGET_TEMPLATE = "HEYUAN_BUSINESS_UNIFIED_SCHEDULE_V1"
ITEM_SHEET = "ITEM表"
ORDER_SHEET = "接单表"
REVIEW_SHEET = "正单评审表"
SHEETS = (ORDER_SHEET, REVIEW_SHEET, ITEM_SHEET)

ORDER_HEADERS = (
    "客出单日期", "订单类型", "Contract No.", "SO#/Reference", "P/O#:",
    "客名", "产品编号", "产品中文名称", "产品英文名称", "數量", "装箱",
    "箱数", "国家标准", "单价HK", "金额HK", "单价USD", "总金额USD",
    "验货日期", "客要求走货期", "备注",
)
ITEM_HEADERS = (
    "客出单日期", "订单类型", "生产单号", "Contract No.", "SO#/Reference",
    "P/O#:", "客名", "产品编号", "产品中文名称", "产品英文名称", "數量",
    "装箱", "箱数", "国家标准", "说明书", "贴纸", "外箱贴纸", "箱唛",
    "布标规格", "包装", "日期码", "是否上系统", "放产日期", "抽板日期",
    "验货日期", "客要求走货期", "第三方公证行验货", "验货结果", "日期码", "备注",
)


class CustomerOrderUnifiedError(CustomerOrderWorkbookError):
    """The uploaded workbook is not the approved Heyuan unified schedule."""


@dataclass
class UnifiedHistoryRow:
    row: int
    source_sheet: str = 'ITEM表'
    received_date: str = ""
    order_type: str = ""
    production_no: str = ""
    contract_no: str = ""
    reference_no: str = ""
    po_no: str = ""
    customer_name: str = ""
    product_no: str = ""
    product_name_zh: str = ""
    product_name_en: str = ""
    quantity: str = ""
    units_per_carton: str = ""
    standard: str = ""
    packaging: str = ""
    date_code: str = ""
    requested_ship_date: str = ""
    unit_price_hkd: str = ""
    extras: dict[str, str] = field(default_factory=dict)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _key(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", _text(value).upper())


def _number(value: Any) -> Decimal | None:
    text = _text(value).replace(",", "")
    if not text:
        return None
    try:
        result = Decimal(text)
    except InvalidOperation:
        return None
    return result if result.is_finite() else None


def _number_text(value: Decimal | None, places: int | None = None) -> str:
    if value is None:
        return ""
    if places is not None:
        value = value.quantize(Decimal(1).scaleb(-places))
    text = format(value, "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def _excel_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    return value


def _date_value(value: Any) -> Any:
    text = _text(value)
    if not text:
        return None
    try:
        return date.fromisoformat(text)
    except ValueError:
        return text


def _combined_hash(file_names: list[str], hashes: list[str]) -> str:
    if len(hashes) == 1:
        return hashes[0]
    payload = "\n".join(
        f"{name}:{digest}" for name, digest in zip(file_names, hashes, strict=True)
    )
    return sha256(payload.encode("utf-8")).hexdigest()


def _load_workbook(content: bytes):
    try:
        return openpyxl.load_workbook(BytesIO(content), data_only=False)
    except Exception as exc:
        raise CustomerOrderUnifiedError(f"无法读取河源业务统一排期：{exc}") from exc


def _normalized_header(value: Any) -> str:
    return re.sub(r"\s+", "", _text(value)).replace("：", ":")


def ensure_unified_schedule(
    schedule_file_name: str, schedule_content: bytes, *, factory_id: str = "", customer_code: str = "",
) -> None:
    if Path(schedule_file_name).suffix.lower() != ".xlsx":
        raise CustomerOrderUnifiedError("客户排期必须使用 .xlsx 格式的《河源业务统一排期》")
    if len(schedule_content) > MAX_SCHEDULE_BYTES:
        raise CustomerOrderUnifiedError("客户排期超过 35MB 限制")
    workbook = _load_workbook(schedule_content)
    try:
        item_names = huaxing.item_sheets(workbook, customer_code) if huaxing.enabled(factory_id, customer_code) else [ITEM_SHEET]
        missing = [name for name in (ORDER_SHEET, REVIEW_SHEET, *item_names) if name not in workbook.sheetnames]
        if missing:
            raise CustomerOrderUnifiedError(
                "排期结构不匹配《河源业务统一排期》，缺少工作表：" + "、".join(missing)
            )
        expected_by_sheet = {
            ORDER_SHEET: ORDER_HEADERS[:19],
            REVIEW_SHEET: ORDER_HEADERS[:19],
            **{name: ITEM_HEADERS[:28] for name in item_names},
        }
        for sheet_name, expected in expected_by_sheet.items():
            worksheet = workbook[sheet_name]
            actual = tuple(
                _normalized_header(worksheet.cell(3, column).value)
                for column in range(1, len(expected) + 1)
            )
            normalized_expected = tuple(_normalized_header(value) for value in expected)
            if actual != normalized_expected:
                first = next(
                    (
                        index + 1
                        for index, (left, right) in enumerate(zip(actual, normalized_expected, strict=True))
                        if left != right
                    ),
                    1,
                )
                raise CustomerOrderUnifiedError(
                    f"{sheet_name} 第3行公共字段不匹配（第{first}列应为“{expected[first - 1]}”）"
                )
    finally:
        workbook.close()


def is_unified_schedule(schedule_content: bytes) -> bool:
    try:
        workbook = openpyxl.load_workbook(BytesIO(schedule_content), read_only=True, data_only=False)
    except Exception:
        return False
    try:
        return set(SHEETS).issubset(workbook.sheetnames) or set((ORDER_SHEET, REVIEW_SHEET, *huaxing.BUZZBEE_SHEETS)).issubset(workbook.sheetnames)
    finally:
        workbook.close()


def _marker_row(worksheet) -> int:
    candidates = [row for (row, _), cell in worksheet._cells.items()
                  if row >= 4 and "取消单" in _text(cell.value)]
    if candidates:
        return min(candidates)
    raise CustomerOrderUnifiedError(f"{worksheet.title} 未找到“取消单”边界")


def read_unified_history(
    schedule_content: bytes, *, factory_id: str = "", customer_code: str = "",
) -> list[UnifiedHistoryRow]:
    workbook = _load_workbook(schedule_content)
    try:
        if huaxing.enabled(factory_id, customer_code):
            return huaxing.read_history(workbook, customer_code)
        huakang = huakang_unified.enabled(factory_id, customer_code)
        regional_mapping = regional.enabled(factory_id, customer_code)
        item_sheet = workbook[ITEM_SHEET]
        order_sheet = workbook[ORDER_SHEET]
        marker = _marker_row(item_sheet)
        price_by_key: dict[tuple[str, str], str] = {}
        for row in range(4, _marker_row(order_sheet)):
            reference = _key(order_sheet.cell(row, 4).value)
            if huakang:
                reference = _key(huakang_unified.history_reference(
                    customer_code, _text(order_sheet.cell(row, 3).value),
                    _text(order_sheet.cell(row, 4).value), _text(order_sheet.cell(row, 5).value),
                ))
            elif regional_mapping:
                reference = _key(regional.history_reference(
                    factory_id, customer_code, _text(order_sheet.cell(row, 3).value),
                    _text(order_sheet.cell(row, 4).value), _text(order_sheet.cell(row, 5).value),
                ))
            product = _key(order_sheet.cell(row, 7).value)
            price = _text(order_sheet.cell(row, 14).value)
            if reference and product and price:
                price_by_key[(reference, product)] = price

        history: list[UnifiedHistoryRow] = []
        extra_columns = huakang_unified.extra_columns(item_sheet, customer_code) if huakang else {}
        if regional_mapping:
            extra_columns = regional.extra_columns(item_sheet, factory_id, customer_code)
        remark_column = next((c.column for c in item_sheet[3] if _normalized_header(c.value) == "备注"), None)
        for row in range(4, (item_sheet.max_row + 1) if huakang or regional_mapping else marker):
            remark = _text(item_sheet.cell(row, remark_column).value) if remark_column else ""
            if "模拟数据" in remark:
                continue
            reference = _text(item_sheet.cell(row, 5).value)
            if huakang:
                reference = huakang_unified.history_reference(
                    customer_code, _text(item_sheet.cell(row, 4).value), reference,
                    _text(item_sheet.cell(row, 6).value),
                )
            elif regional_mapping:
                reference = regional.history_reference(
                    factory_id, customer_code, _text(item_sheet.cell(row, 4).value),
                    reference, _text(item_sheet.cell(row, 6).value),
                )
            product = _text(item_sheet.cell(row, 8).value)
            if not reference or not product:
                continue
            history.append(UnifiedHistoryRow(
                row=row,
                received_date=_text(item_sheet.cell(row, 1).value),
                order_type=_text(item_sheet.cell(row, 2).value),
                production_no=_text(item_sheet.cell(row, 3).value),
                contract_no=_text(item_sheet.cell(row, 4).value),
                reference_no=reference,
                po_no=_text(item_sheet.cell(row, 6).value),
                customer_name=_text(item_sheet.cell(row, 7).value),
                product_no=product,
                product_name_zh=_text(item_sheet.cell(row, 9).value),
                product_name_en=_text(item_sheet.cell(row, 10).value),
                quantity=regional.history_number(item_sheet.cell(row, 11), workbook, customer_code) if regional_mapping else _text(item_sheet.cell(row, 11).value),
                units_per_carton=regional.history_number(item_sheet.cell(row, 12), workbook, customer_code, pack=True) if regional_mapping else _text(item_sheet.cell(row, 12).value),
                standard=_text(item_sheet.cell(row, 14).value),
                packaging=_text(item_sheet.cell(row, 20).value),
                date_code=_text(item_sheet.cell(row, 21).value) or (
                    _text(item_sheet.cell(row, 29).value)
                    if _normalized_header(item_sheet.cell(3, 29).value) == "日期码" else ""
                ),
                requested_ship_date=_text(item_sheet.cell(row, 26).value),
                unit_price_hkd=price_by_key.get((_key(reference), _key(product)), ""),
                extras={
                    field: _text(item_sheet.cell(row, column).value)
                    for field, column in extra_columns.items()
                    if item_sheet.cell(row, column).data_type != "f"
                },
            ))
        return history
    finally:
        workbook.close()


def _unique_history_value(
    history: list[UnifiedHistoryRow], product_no: str, field: str, customer_name: str = "",
) -> str:
    product_key = _key(product_no)
    candidates = [row for row in history if _key(row.product_no) == product_key]
    customer_key = _key(customer_name)
    customer_matches = [row for row in candidates if customer_key and _key(row.customer_name) == customer_key]
    if customer_matches:
        candidates = customer_matches
    values = {_text(getattr(row, field)) for row in candidates if _text(getattr(row, field))}
    return next(iter(values)) if len(values) == 1 else ""


def _issue(
    severity: str, code: str, field: str, message: str, *, can_skip: bool = False,
    skip_label: str = "",
) -> dict[str, Any]:
    return {
        "severity": severity,
        "code": code,
        "field": field,
        "message": message,
        "can_skip": can_skip,
        "skip_key": "",
        "skip_label": skip_label,
    }


def _refresh_status(row: dict[str, Any]) -> None:
    status = (
        "blocked" if any(issue.get("severity") == "blocked" for issue in row["issues"])
        else "warning" if row["issues"] else "valid"
    )
    row["status"] = status
    row["status_label"] = {"valid": "可导出", "warning": "需复核", "blocked": "已阻断"}[status]


def _attach_available_fields(row: dict[str, Any], *sources: dict[str, Any]) -> dict[str, Any]:
    aliases = {
        "order_type": ("order_type", "document_type"),
        "production_no": ("production_no", "huaxing_po"),
        "manual": ("manual", "instructions", "instruction"),
        "label": ("label", "sticker"),
        "customer_label": ("customer_label", "outer_carton_sticker"),
        "carton_mark": ("carton_mark", "shipping_mark"),
        "fabric_label": ("fabric_label",),
        "date_code": ("date_code", "expected_date_code", "calculated_date_code"),
        "system_status": ("system_status", "on_system"),
        "release_date": ("release_date", "production_release_date"),
        "sample_date": ("sample_date",),
        "third_party_inspection": ("third_party_inspection", "inspection_company"),
        "inspection_result": ("inspection_result",),
    }
    for target, candidates in aliases.items():
        if _text(row.get(target)):
            continue
        for source in sources:
            value = next((_text(source.get(name)) for name in candidates if _text(source.get(name))), "")
            if value:
                row[target] = value
                break
    return row


def _enrich_row(
    row: dict[str, Any], history: list[UnifiedHistoryRow], customer_name: str,
) -> dict[str, Any]:
    row.setdefault("row_role", "detail")
    row.setdefault("parent_product_no", "")
    row.setdefault("issues", [])
    row["target_template"] = huaxing.TARGET_TEMPLATE if row.get('_huaxing_customer') else TARGET_TEMPLATE
    row.setdefault("item_sheet_name", ITEM_SHEET)
    row["customer_name"] = _text(row.get("customer_name")) or customer_name
    row["customer_country"] = _text(row.get("customer_country")) or " / ".join(
        value for value in (row["customer_name"], _text(row.get("country"))) if value
    )
    product_no = _text(row.get("product_no"))
    inherited = {
        "product_name_zh": "product_name_zh",
        "product_name_en": "product_name_en",
        "units_per_carton": "units_per_carton",
        "standard": "standard",
        "packaging": "packaging",
        "unit_price_hkd": "unit_price_hkd",
    }
    for target, source in inherited.items():
        if (row.get('_regional_customer') == 'ubtech' or row.get('_jakks_supplementary')) and target == 'unit_price_hkd':
            continue
        if not _text(row.get(target)):
            value = _unique_history_value(history, product_no, source, row["customer_name"])
            if value:
                row[target] = value
                row.setdefault("lineage", {})[target] = (
                    f"河源业务统一排期 · 货号 {product_no} 唯一历史值"
                )

    quantity = _number(row.get("quantity"))
    units = _number(row.get("units_per_carton"))
    price = _number(row.get("unit_price_hkd"))
    if quantity is not None and units not in (None, Decimal(0)):
        row["carton_count"] = _number_text(quantity / units)
    if quantity is not None and price is not None:
        row["amount_hkd"] = _number_text(quantity * price, 2)

    reference_no = _text(row.get("reference_no")) or _text(row.get("contract_no")) or _text(row.get("po_no"))
    row["reference_no"] = reference_no
    cleaned: list[dict[str, Any]] = []
    for existing_issue in row["issues"]:
        code = _text(existing_issue.get("code"))
        # Unified ITEM columns are a writable superset, not a mandatory field set.
        # Customer-specific parsers were built for older schedules and may flag
        # absent optional columns as blockers.  Keep those values blank here; the
        # unified writer only requires the two keys needed to place and dedupe a
        # line reliably.
        if code.startswith("missing_") or code == "product_not_in_schedule":
            continue
        cleaned.append(existing_issue)
    row["issues"] = cleaned

    required = (
        ("contract_no", reference_no, "缺少 Contract No. / SO#/Reference"),
        ("product_no", product_no, "缺少产品编号"),
    )
    existing_fields = {
        _text(issue.get("field")) for issue in row["issues"] if issue.get("severity") == "blocked"
    }
    for field, value, message in required:
        if not value and field not in existing_fields:
            row["issues"].append(_issue("blocked", f"missing_{field}", field, message))

    huakang_customer = row.get("_huakang_customer", "")
    identity = huakang_unified.identity(huakang_customer, reference_no, product_no)
    matches = [
        item for item in history
        if identity[0] and identity[1]
        and huakang_unified.identity(huakang_customer, item.reference_no, item.product_no) == identity
    ]
    if row.get('_huaxing_customer') == 'edu':
        reference_matches = matches
        matches = huaxing.edu_history_matches(row, [item for item in history if _key(item.product_no) == _key(product_no)])
        if any(item not in matches for item in reference_matches):
            row['issues'].append(_issue('blocked', 'edu_reference_collision', 'contract_no',
                'EDU 内部编号与产品已被另一客户合同/PO占用，请核对原订单；不能作为普通重复单放行'))
    if matches and not any(
        issue.get("code") in {"duplicate_reference", "existing_order_line", "duplicate_existing_order"}
        for issue in row["issues"]
    ):
        existing_quantities = {_number(item.quantity) for item in matches}
        same_quantity = quantity is not None and quantity in existing_quantities
        row["issues"].append(_issue(
            "blocked",
            "duplicate_existing_order" if same_quantity else "existing_quantity_conflict",
            "contract_no" if same_quantity else "quantity",
            (
                f"统一排期 {matches[0].source_sheet}第 {matches[0].row} 行已有相同 {'客户合同/PO' if row.get('_huaxing_customer') == 'edu' else 'SO#/Reference'} 与产品"
                if same_quantity
                else f"统一排期已有相同 SO#/Reference 与产品，但数量不同"
            ),
            can_skip=same_quantity,
            skip_label="确认重复导入统一排期已有订单" if same_quantity else "",
        ))
    for issue in row["issues"]:
        if issue.get("can_skip") and not issue.get("skip_key"):
            issue["skip_key"] = f"{row['id']}|{issue.get('code', 'issue')}|{issue.get('field', 'row')}"
    _refresh_status(row)
    return row


def _mark_batch_duplicates(rows: list[dict[str, Any]]) -> None:
    seen: dict[tuple[str, ...], dict[str, Any]] = {}
    edu_references: dict[tuple[str, str], tuple[str, ...]] = {}
    for row in rows:
        if row.get("row_role") == "parent":
            continue
        identity = huakang_unified.identity(row.get("_huakang_customer", ""), row.get("reference_no"), row.get("product_no"))
        is_edu = row.get('_huaxing_customer') == 'edu'
        if is_edu:
            reference_identity = identity
            identity = huaxing.edu_identity(row.get('contract_no'), row.get('po_no'), row.get('product_no'))
            previous_identity = edu_references.setdefault(reference_identity, identity)
            if all(reference_identity) and previous_identity != identity:
                row['issues'].append(_issue('blocked', 'edu_reference_collision', 'contract_no',
                    '本批不同客户合同/PO使用了相同 EDU 内部编号与产品，请核对原订单'))
                _refresh_status(row)
        if not identity[-1] or not any(identity[:-1]):
            continue
        previous = seen.get(identity)
        if previous is None:
            seen[identity] = row
            continue
        conflict = is_edu and _number(row.get('quantity')) != _number(previous.get('quantity'))
        code = 'existing_quantity_conflict' if conflict else 'duplicate_batch_order_line'
        field = 'quantity' if conflict else 'contract_no'
        label = '客户合同/PO 与产品' if is_edu else 'SO#/Reference 与产品'
        issue = _issue(
            "blocked", code, field,
            f"本批与 {previous.get('source_po_file_name') or '之前文件'} 存在相同 {label}" + ('，但数量不同，按修改/补单阻断' if conflict else ''),
            can_skip=not conflict, skip_label="确认本批重复订单行仍需分别写入" if not conflict else '',
        )
        if not conflict:
            issue["skip_key"] = f"{row['id']}|{code}|{field}"
        row["issues"].append(issue)
        _refresh_status(row)


def _validate_po_batch(po_files: list[tuple[str, bytes]]) -> None:
    if not po_files:
        raise CustomerOrderUnifiedError("请至少上传一份 PO 文件")
    if len(po_files) > MAX_BATCH_PO_FILES:
        raise CustomerOrderUnifiedError(f"单批最多上传 {MAX_BATCH_PO_FILES} 份 PO 文件")
    if sum(len(content) for _, content in po_files) > MAX_BATCH_PO_BYTES:
        raise CustomerOrderUnifiedError("本批 PO 文件合计超过 80MB 限制")
    oversized = [name for name, content in po_files if len(content) > MAX_PO_BYTES]
    if oversized:
        raise CustomerOrderUnifiedError("PO 文件超过 12MB 限制：" + "、".join(oversized))


def _factory_allowed(customer_code: str, factory_id: str) -> bool:
    if customer_code == 'ubtech':
        return factory_id == 'huakang-d'
    if customer_code == 'disney':
        return factory_id in {'huaxing', 'huakang-d'}
    if customer_code == 'spin-master' or (customer_code == 'jp' and factory_id == 'huakang-d'):
        return False
    if customer_code in {"buzzbee", "dickie", "caixing", "disney", "edu", "360", "yinhui", "seasons", "maxx", "shushupapa", "barter"}:
        return factory_id == "huaxing" or (customer_code == "360" and factory_id == "huakang-a") or (customer_code == "maxx" and factory_id in {"huakang-c", "huakang-d"})
    if customer_code in {"casdon", "jakks", "simba", "spin", "spin-master", "goliath"}:
        return factory_id == "huadeng"
    if customer_code in {"green-toys", "headstart"}:
        return factory_id == "huakang-a"
    if customer_code in {"index", "jazwares", "strottman", "jp"}:
        return factory_id in {"huakang-c", "huakang-d"}
    return False


def _customer_name(customer_code: str, factory_id: str) -> str:
    if customer_code == 'ubtech':
        return '优必选'
    if customer_code == "360" and factory_id == "huakang-a":
        from app.services.customer_order_huakang_a import get_huakang_a_customer_mapping
        return get_huakang_a_customer_mapping(customer_code).name
    if customer_code == "maxx" and factory_id in {"huakang-c", "huakang-d"}:
        from app.services.customer_order_huakang_c import get_huakang_c_customer_mapping
        return get_huakang_c_customer_mapping(customer_code).name
    if customer_code in {"buzzbee", "dickie", "caixing"}:
        return {"buzzbee": "BuzzBee", "dickie": "Dickie", "caixing": "彩星"}[customer_code]
    if customer_code in {"disney", "edu", "360", "yinhui", "seasons", "maxx", "shushupapa", "barter"}:
        from app.services.customer_order_huaxing import get_huaxing_customer_mapping
        return get_huaxing_customer_mapping(customer_code).name
    if customer_code in {"casdon", "jakks", "simba", "spin", "spin-master", "goliath"}:
        from app.services.customer_order_huadeng import get_huadeng_customer_mapping
        return get_huadeng_customer_mapping(customer_code).name
    if customer_code in {"green-toys", "headstart"}:
        from app.services.customer_order_huakang_a import get_huakang_a_customer_mapping
        return get_huakang_a_customer_mapping(customer_code).name
    from app.services.customer_order_huakang_c import get_huakang_c_customer_mapping
    return get_huakang_c_customer_mapping(customer_code).name


def _parse_special_rows(
    customer_code: str,
    po_files: list[tuple[str, bytes]],
    received_date: str,
    history: list[UnifiedHistoryRow],
) -> tuple[list[dict[str, Any]], list[str], str]:
    if customer_code == "buzzbee":
        from app.services import customer_order_buzzbee as service

        schedule_index: dict[str, Any] = {}
        for item in history:
            product_key = service._clean_identifier(item.product_no)
            if not product_key:
                continue
            schedule_index[product_key] = service.ScheduleLookup(
                product_name_zh=_unique_history_value(history, item.product_no, "product_name_zh", item.customer_name),
                unit_price_hkd=_number(_unique_history_value(history, item.product_no, "unit_price_hkd", item.customer_name)),
                existing_rows=[],
            )
        parsed_lines = []
        for file_name, content in po_files:
            parsed = service.parse_po(file_name, content)
            for line in parsed:
                line.source_file_name = file_name
            parsed_lines.extend(parsed)
        return (
            service._build_preview_rows(parsed_lines, schedule_index, received_date),
            [],
            "BUZZBEE_PO_AUTO_DETECT_V1",
        )

    if customer_code == "dickie":
        from app.services import customer_order_dickie as service

        product_index: dict[str, Any] = {}
        for item in history:
            product_key = service._product_key(item.product_no)
            if not product_key:
                continue
            lookup = product_index.setdefault(product_key, service.DickieProductLookup())
            lookup.product_name_zh = _unique_history_value(history, item.product_no, "product_name_zh", item.customer_name)
            lookup.customer_name = _unique_history_value(history, item.product_no, "customer_name", item.customer_name)
            lookup.packaging = _unique_history_value(history, item.product_no, "packaging", item.customer_name)
        rows: list[dict[str, Any]] = []
        for file_name, content in po_files:
            parsed_orders = service.parse_dickie_pdf_orders(
                file_name, content, fallback_received_date=received_date,
            )
            for parsed in parsed_orders:
                row = service._preview_row(
                    parsed,
                    service._lookup_product(product_index, parsed.product_no),
                    product_index,
                    file_name=file_name,
                    row_index=len(rows) + 1,
                )
                row["received_date"] = received_date
                rows.append(row)
        return rows, ["Dickie 扫描 PDF 的 OCR 字段仍需在预览中人工核对。"], service.INPUT_TEMPLATE

    if customer_code == "caixing":
        from app.services import customer_order_caixing as service

        class UnifiedMatrix:
            @staticmethod
            def matrix_columns_for(_product_no: str) -> tuple[str, str]:
                return "", ""

        rows = []
        for file_name, content in po_files:
            for parsed in service.parse_caixing_pdf(file_name, content):
                rows.append(service._preview_row(
                    parsed,
                    file_name=file_name,
                    row_index=len(rows) + 1,
                    received_date=received_date,
                    schedule=UnifiedMatrix(),
                    existing=set(),
                ))
        return rows, ["彩星套装仅写入小货号明细，避免大货号汇总与明细重复计数。"], service.INPUT_TEMPLATE

    raise CustomerOrderUnifiedError(f"不支持的客户映射：{customer_code}")


def _disney_schedule_stub(history: list[UnifiedHistoryRow]) -> dict[str, Any]:
    from app.services.huaxing_order_legacy import disney_schedule

    masters: dict[str, dict[str, Any]] = {}
    for item in history:
        key = disney_schedule._key(item.product_no)
        if not key:
            continue
        masters[key] = {
            "product_name_zh": _unique_history_value(history, item.product_no, "product_name_zh", item.customer_name),
            "product_name_en": _unique_history_value(history, item.product_no, "product_name_en", item.customer_name),
            "name_conflict": False,
        }
    return {"records": [], "item_master": masters, "sheet": ITEM_SHEET}


def _parse_huaxing_rows(
    customer_code: str,
    po_files: list[tuple[str, bytes]],
    received_date: str,
    history: list[UnifiedHistoryRow],
) -> tuple[list[dict[str, Any]], list[str], str]:
    from app.services import customer_order_huaxing as service
    from app.services.huaxing_order_legacy import (
        disney_schedule,
        edu_po_parser,
        edu_schedule,
        multi_po_parser,
        multi_schedule,
        shixin_schedule,
        three_sixty_po_parser,
        three_sixty_schedule,
        yinhui_po_parser,
        yinhui_schedule,
    )

    spec = service.get_huaxing_customer_mapping(customer_code)
    warnings: list[str] = []
    records: list[dict[str, Any]] = []

    if customer_code == "edu":
        parsed_files: list[dict[str, Any]] = []
        next_number = service._next_edu_number([
            {"huaxing_po": item.reference_no} for item in history
        ])
        for file_name, content in po_files:
            parsed = edu_po_parser.parse_po_file(content, filename=file_name)
            if parsed.get("customer_type") != "EDU":
                raise CustomerOrderUnifiedError(
                    f"{file_name}：识别为 {parsed.get('customer_type') or '未知客户'}，当前入口只处理 EDU"
                )
            for record in parsed.get("rows", []):
                record["_source_po_file_name"] = file_name
                edu_schedule.add_derived_fields(record)
            warnings.extend(f"{file_name}：{message}" for message in parsed.get("warnings", []))
            parsed_files.append(parsed)
        records, dedupe_warnings = edu_po_parser.merge_edu_po_results(parsed_files)
        huaxing.allocate_edu_references(records, next_number)
        warnings.extend(dedupe_warnings)

    elif customer_code == "360":
        parsed_files = []
        for file_name, content in po_files:
            parsed = three_sixty_po_parser.parse_po_file(content, filename=file_name)
            for record in parsed.get("rows", []):
                record["_source_po_file_name"] = file_name
            warnings.extend(f"{file_name}：{message}" for message in parsed.get("warnings", []))
            parsed_files.append(parsed)
        records = [
            three_sixty_schedule.add_derived_fields(record)
            for record in three_sixty_po_parser.merge_po_results(parsed_files)
        ]

    elif customer_code == "yinhui":
        seen: set[tuple[str, ...]] = set()
        for file_name, content in po_files:
            parsed = yinhui_po_parser.parse_po(content, file_name)
            warnings.extend(f"{file_name}：{message}" for message in parsed.get("warnings", []))
            for record in parsed.get("rows", []):
                identity = tuple(_key(record.get(field)) for field in (
                    "contract_no", "item_no", "quantity", "po_ship_date", "unit_price_usd", "total_usd",
                ))
                if identity in seen:
                    continue
                seen.add(identity)
                record["_source_po_file_name"] = file_name
                yinhui_schedule.add_derived_fields(record)
                records.append(record)

    elif customer_code == "disney":
        documents = []
        for file_name, content in po_files:
            document = disney_schedule.parse_po(content, file_name)
            for record in document.get("rows", []):
                record["_source_po_file_name"] = file_name
            warnings.extend(f"{file_name}：{message}" for message in document.get("warnings", []))
            documents.append(document)
        records, revision_warnings = disney_schedule.merge_revisions(documents)
        warnings.extend(revision_warnings)
        schedule_stub = _disney_schedule_stub(history)
        records = [disney_schedule.add_schedule_fields(record, schedule_stub) for record in records]

    elif customer_code == "seasons":
        documents = []
        for file_name, content in po_files:
            document = shixin_schedule.parse_order_documents(content, file_name, [])
            for record in document.get("records", []):
                record["_source_po_file_name"] = file_name
            warnings.extend(document.get("warnings", []))
            documents.append(document)
        reconciled = shixin_schedule.reconcile_documents(documents, [])
        records = reconciled.get("records", [])
        warnings.extend(reconciled.get("warnings", []))

    else:
        orders: list[dict[str, Any]] = []
        for file_name, content in po_files:
            order = multi_po_parser.parse_po(content, file_name)
            detected = _text(order.get("client")).lower()
            if detected != customer_code:
                raise CustomerOrderUnifiedError(
                    f"{file_name}：识别为 {order.get('client_name') or '未知客户'}，当前入口只处理 {spec.name}"
                )
            order["filename"] = file_name
            warnings.extend(f"{file_name}：{message}" for message in order.get("warnings", []))
            orders.append(order)
        orders, revision_warnings = service._dedupe_multi_orders(orders)
        warnings.extend(revision_warnings)
        today = date.fromisoformat(received_date)
        for order in orders:
            if customer_code == "barter":
                order["po_date"] = received_date
            for line in order.get("lines", []):
                if line.get("is_charge"):
                    continue
                record = multi_schedule._line_values(order, line)
                record["_source_po_file_name"] = _text(order.get("filename"))
                record["source_page"] = line.get("source_page")
                record["flags"] = multi_schedule._record_flags(record, customer_code, today)
                records.append(record)

    if not records:
        raise CustomerOrderUnifiedError(f"本批文件未识别到可生成的 {spec.name} 新单明细")
    rows = []
    for index, record in enumerate(records, start=1):
        row = service._preview_row(
            customer_code=customer_code,
            spec=spec,
            received_date=received_date,
            record=record,
            index=index,
            default_source_file=po_files[0][0],
            sheet_name=ITEM_SHEET,
        )
        rows.append(huaxing.map_record(_attach_available_fields(row, record), record, customer_code))
    return rows, warnings, spec.input_template


def _parse_huadeng_rows(
    customer_code: str,
    po_files: list[tuple[str, bytes]],
    received_date: str,
    history: list[UnifiedHistoryRow],
) -> tuple[list[dict[str, Any]], list[str], str]:
    from app.services import customer_order_huadeng as service
    from app.services.huadeng_order_legacy import (
        casdon_po_parser,
        casdon_schedule,
        jakks_po_parser,
        simba_po_parser,
        simba_schedule,
        spin_master_new_order_writer,
        spin_master_parser,
        spin_po_parser,
        spin_schedule,
    )

    spec = service.get_huadeng_customer_mapping(customer_code)
    warnings: list[str] = []
    records: list[dict[str, Any]] = []

    if customer_code == 'goliath':
        from app.services.huadeng_order_legacy.goliath_po_parser import parse_po
        seen_files: set[str] = set()
        for file_name, content in po_files:
            digest = sha256(content).hexdigest()
            if digest in seen_files:
                warnings.append(f'{file_name}：完全相同的重复文件未再次读取。')
                continue
            seen_files.add(digest)
            records.extend(parse_po(content, file_name))
    elif customer_code in {"casdon", "spin"}:
        is_casdon = customer_code == "casdon"
        parser_class = casdon_po_parser.CasdonPOParser if is_casdon else spin_po_parser.SpinPOParser
        validator = casdon_po_parser.validate if is_casdon else spin_po_parser.validate
        schedule_module = casdon_schedule if is_casdon else spin_schedule
        index = schedule_module.MasterIndex()
        orders: list[dict[str, Any]] = []
        with TemporaryDirectory(prefix=f"huadeng-unified-{customer_code}-") as temp_dir:
            root = Path(temp_dir)
            for number, (file_name, content) in enumerate(po_files, start=1):
                safe_name = re.sub(r"[\\/:*?\"<>|]+", "_", Path(file_name).name) or "po"
                po_path = root / f"{number:03d}-{safe_name}"
                po_path.write_bytes(content)
                try:
                    order = parser_class().parse(str(po_path))
                except Exception as exc:
                    warnings.append(f"{file_name}：解析失败（{exc}），本文件未进入新单")
                    continue
                order["filename"] = file_name
                if is_casdon:
                    order["email_received_date"] = received_date
                warnings.extend(validator(order, file_name))
                orders.append(order)
        orders, dedupe_warnings = service._dedupe_revision_orders(
            orders,
            quality_fields=(
                "po_number", "po_date", "ship_date", "your_reference", "customer", "customer_po_header",
            ) if is_casdon else (),
        )
        warnings.extend(dedupe_warnings)
        for order in orders:
            for line in order.get("lines", []):
                values = schedule_module._compose_row(order, line, index)
                record = service._legacy_row_record(customer_code, order, line, values)
                if not is_casdon:
                    record['unified_reference'] = values.get(schedule_module.COL['item'])
                    record['material_group'] = line.get('material_group')
                for field_name in ('carton_mark', 'customer_label', 'date_code', 'container_type', 'brand', 'certificate'):
                    column = schedule_module.COL.get(field_name)
                    if column:
                        record[field_name] = values.get(column)
                records.append(record)

    elif customer_code == "jakks":
        dataset = {
            "records": [
                {
                    "sku": item.product_no,
                    "product": item.product_name_zh,
                    "date_code": item.date_code,
                    "contract_no": item.contract_no,
                    "customer_po": item.po_no,
                    "contact": "",
                    "quantity": item.quantity,
                }
                for item in history
            ]
        }
        parsed_orders: list[dict[str, Any]] = []
        with TemporaryDirectory(prefix="huadeng-unified-jakks-") as temp_dir:
            root = Path(temp_dir)
            for number, (file_name, content) in enumerate(po_files, start=1):
                if re.search(r"(?:^|[-_])CXL(?:[-_]|$)", Path(file_name).stem, re.I):
                    warnings.append(f"{file_name}：文件名含 CXL，按取消单拦截。")
                    continue
                po_path = root / f"{number:03d}_{Path(file_name).name}"
                po_path.write_bytes(content)
                try:
                    order = jakks_po_parser.parse_po(po_path)
                except Exception as exc:
                    warnings.append(f"{file_name}：{exc}")
                    continue
                order["filename"] = file_name
                service._apply_jakks_schedule_data(order, dataset)
                parsed_orders.append(order)
        seen_exact: set[tuple[str, ...]] = set()
        seen_business: set[tuple[str, ...]] = set()
        for order in parsed_orders:
            warnings.extend(f"{order.get('filename')}：{message}" for message in order.get("warnings", []))
            for raw_line in order.get("lines", []):
                line = dict(raw_line)
                for field_name in (
                    "order_date", "contact", "customer_po", "confirmation_no", "contract_no",
                    "customer", "country", "ship_date", "national_standard",
                ):
                    if not line.get(field_name):
                        line[field_name] = order.get(field_name)
                line["source_file"] = order.get("filename") or ""
                exact = service._jakks_line_identity(line)
                business = service._jakks_line_identity(line, business=True)
                if exact in seen_exact or business in seen_business:
                    warnings.append(f"{line['source_file']} / {line.get('item_no')}：批内重复明细未再次写入。")
                    continue
                seen_exact.add(exact)
                seen_business.add(business)
                records.append(line)

    elif customer_code == "simba":
        excel_stems = {
            Path(name).stem.casefold() for name, _ in po_files
            if Path(name).suffix.lower() in {".xlsx", ".xlsm"}
        }
        selected: list[tuple[str, bytes]] = []
        seen_hashes: set[str] = set()
        for file_name, content in po_files:
            digest = sha256(content).hexdigest()
            if digest in seen_hashes:
                warnings.append(f"{file_name}：完全相同的重复文件未再次处理。")
                continue
            seen_hashes.add(digest)
            if Path(file_name).suffix.lower() == ".pdf" and Path(file_name).stem.casefold() in excel_stems:
                warnings.append(f"{file_name}：检测到同名 Excel，已优先采用 Excel。")
                continue
            selected.append((file_name, content))
        merged: dict[tuple[str, ...], dict[str, Any]] = {}
        order_keys: list[tuple[str, ...]] = []
        for file_name, content in selected:
            try:
                parsed = simba_po_parser.parse_po_file(content, filename=file_name)
            except Exception as exc:
                warnings.append(f"{file_name}：{exc}")
                continue
            warnings.extend(f"{file_name}：{message}" for message in parsed.get("warnings", []))
            for source in parsed.get("rows", []):
                record = dict(source)
                for field_name in ("outer_length_cm", "outer_width_cm", "outer_height_cm"):
                    record.pop(field_name, None)
                record["english_name"] = service._simba_first_english_sentence(record.get("english_name"))
                record["source_file"] = file_name
                record["_source_po_file_name"] = file_name
                record["order_date"] = received_date
                simba_schedule.add_derived_fields(record)
                key = service._simba_identity(record)
                current = merged.get(key)
                if current is None:
                    merged[key] = record
                    order_keys.append(key)
                elif service._simba_signature(current) != service._simba_signature(record) and service._simba_priority(file_name, record) > service._simba_priority(_text(current.get("source_file")), current):
                    merged[key] = record
                else:
                    warnings.append(f"{file_name} / {record.get('item_no')}：同单重复或低优先级版本未再次写入。")
        records = [merged[key] for key in order_keys]

    else:
        seen: set[tuple[str, ...]] = set()
        with TemporaryDirectory(prefix="huadeng-unified-spin-master-") as temp_dir:
            root = Path(temp_dir)
            for number, (file_name, content) in enumerate(po_files, start=1):
                po_path = root / f"{number:03d}{Path(file_name).suffix.lower()}"
                po_path.write_bytes(content)
                try:
                    parsed = spin_master_parser.parse_po_file(po_path)
                except Exception as exc:
                    warnings.append(f"{file_name}：{exc}")
                    continue
                parsed["source"] = file_name
                for record, source_item in zip(spin_master_new_order_writer._records(parsed), parsed.get('items', []), strict=True):
                    record['material_group'] = source_item.get('material_group')
                    record['unified_reference'] = record.get('item_no')
                    record["_source_po_file_name"] = file_name
                    identity = tuple(_key(record.get(field)) for field in (
                        "contract_no", "customer_po", "item_no", "quantity", "ship_date", "unit_price_usd", "total_usd",
                    ))
                    if identity in seen:
                        continue
                    seen.add(identity)
                    records.append(record)

    if not records:
        detail = "；".join(dict.fromkeys(warnings))
        raise CustomerOrderUnifiedError(
            f"本批文件未识别到可生成的 {spec.name} 新单明细" + (f"：{detail}" if detail else "")
        )
    rows = []
    for index, record in enumerate(records, start=1):
        row = service._preview_row(
            customer_code=customer_code,
            spec=spec,
            received_date=received_date,
            record=record,
            index=index,
            default_source_file=po_files[0][0],
            sheet_name=ITEM_SHEET,
        )
        rows.append(regional.map_huadeng(_attach_available_fields(row, record), record, customer_code))
    return rows, warnings, spec.input_template


def _parse_huakang_a_rows(
    customer_code: str,
    po_files: list[tuple[str, bytes]],
    received_date: str,
    history: list[UnifiedHistoryRow],
) -> tuple[list[dict[str, Any]], list[str], str]:
    from app.services import customer_order_huakang_a as service
    from app.services.huakang_a_order_legacy import green_toys_headstart, schedule_parser

    spec = service.get_huakang_a_customer_mapping(customer_code)
    records: list[dict[str, Any]] = []
    warnings: list[str] = []
    if customer_code == "360":
        with TemporaryDirectory(prefix="huakang-a-unified-360-") as temp_dir:
            root = Path(temp_dir)
            for number, (file_name, content) in enumerate(po_files, start=1):
                po_path = root / f"{number:03d}_{Path(file_name).name}"
                po_path.write_bytes(content)
                try:
                    record = schedule_parser.parse_po_file(po_path, root)
                except Exception as exc:
                    raise CustomerOrderUnifiedError(f"{file_name}：无法读取华康A 360 PO：{exc}") from exc
                record["file_name"] = file_name
                record["relative_path"] = file_name
                record["_source_po_file_name"] = file_name
                warnings.extend(f"{file_name}：{message}" for message in record.get("parse_warnings", []))
                records.append(record)
    else:
        parser: Callable[[str, bytes], list[dict[str, Any]]] = (
            green_toys_headstart.parse_green_toys_po
            if customer_code == "green-toys"
            else green_toys_headstart.parse_headstart_pdf
        )
        for file_name, content in po_files:
            try:
                parsed = parser(file_name, content)
            except Exception as exc:
                raise CustomerOrderUnifiedError(f"{file_name}：无法读取 {spec.name} PO：{exc}") from exc
            for record in parsed:
                record["_source_po_file_name"] = file_name
                records.append(record)
    if not records:
        raise CustomerOrderUnifiedError(f"本批文件未识别到可生成的 {spec.name} 新单明细")
    rows = []
    assigned_green_po: dict[tuple[str, str], str] = {}
    for index, record in enumerate(records, start=1):
        if customer_code == "green-toys":
            item_rows: dict[str, list[dict[str, Any]]] = {}
            for item in history:
                item_rows.setdefault(green_toys_headstart._item_key(item.product_no), []).append({
                    "item_no": item.product_no, "description": item.product_name_en,
                })
            record["item_no"] = green_toys_headstart._reconcile_green_item(record, item_rows)
            huakang_unified.reconcile_green_po(record, history, assigned_green_po)
        row = service._preview_row(
            spec=spec,
            received_date=received_date,
            record=record,
            index=index,
            sheet_name=ITEM_SHEET,
        )
        rows.append(huakang_unified.map_record(_attach_available_fields(row, record), record, customer_code))
    return rows, warnings, spec.input_template


def _parse_huakang_c_rows(
    customer_code: str,
    po_files: list[tuple[str, bytes]],
    received_date: str,
    factory_id: str = 'huakang-c',
) -> tuple[list[dict[str, Any]], list[str], str]:
    from app.services import customer_order_huakang_c as service
    from app.services.huakang_c_order_legacy import huakang_po_parser

    spec = service.get_huakang_c_customer_mapping(customer_code)
    parser = huakang_po_parser.HuakangPOParser()
    orders: list[dict[str, Any]] = []
    errors: list[str] = []
    with TemporaryDirectory(prefix=f"huakang-unified-{customer_code}-") as temp_dir:
        root = Path(temp_dir)
        for number, (file_name, content) in enumerate(po_files, start=1):
            po_path = root / f"{number:03d}_{Path(file_name).name}"
            po_path.write_bytes(content)
            try:
                order = parser.parse(str(po_path))
            except Exception as exc:
                errors.append(f"{file_name}：无法读取PO：{exc}")
                continue
            detected = _text(order.get("customer_code")).lower()
            if detected != spec.legacy_code:
                raise CustomerOrderUnifiedError(
                    f"{file_name}：识别为 {detected or '未知客户'}，当前入口只处理 {spec.name}"
                )
            order["filename"] = file_name
            order["po_date"] = received_date
            orders.append(order)
    if not orders:
        raise CustomerOrderUnifiedError("所有文件均解析失败" + (f"：{'；'.join(errors)}" if errors else ""))
    orders, dedupe_warnings = service._deduplicate(orders)
    warnings = [*errors, *dedupe_warnings]
    entries = [(order, line) for order in orders for line in order.get("lines", [])]
    rows = []
    for index, (order, line) in enumerate(entries, start=1):
        row = service._preview_row(
            spec=spec,
            received_date=received_date,
            order=order,
            line=line,
            index=index,
            sheet_name=ITEM_SHEET,
        )
        _attach_available_fields(row, line, order)
        if factory_id == 'huakang-d':
            regional.map_huakang_d(row, order, line, customer_code)
        rows.append(row)
    if not rows:
        raise CustomerOrderUnifiedError(f"本批文件未识别到可生成的 {spec.name} 新单明细")
    return rows, warnings, spec.input_template


def _create_customer_rows(
    customer_code: str,
    factory_id: str,
    po_files: list[tuple[str, bytes]],
    received_date: str,
    history: list[UnifiedHistoryRow],
) -> tuple[list[dict[str, Any]], list[str], str]:
    if customer_code == 'ubtech':
        from app.services.customer_order_ubtech import create_rows
        return create_rows(po_files, received_date, history)
    if customer_code in {"buzzbee", "dickie", "caixing"}:
        return _parse_special_rows(customer_code, po_files, received_date, history)
    if customer_code == "360" and factory_id == "huakang-a":
        return _parse_huakang_a_rows(customer_code, po_files, received_date, history)
    if customer_code == "maxx" and factory_id in {"huakang-c", "huakang-d"}:
        return _parse_huakang_c_rows(customer_code, po_files, received_date, factory_id)
    if customer_code in {"disney", "edu", "360", "yinhui", "seasons", "maxx", "shushupapa", "barter"}:
        return _parse_huaxing_rows(customer_code, po_files, received_date, history)
    if customer_code in {"casdon", "jakks", "simba", "spin", "spin-master", "goliath"}:
        return _parse_huadeng_rows(customer_code, po_files, received_date, history)
    if customer_code in {"green-toys", "headstart"}:
        return _parse_huakang_a_rows(customer_code, po_files, received_date, history)
    return _parse_huakang_c_rows(customer_code, po_files, received_date, factory_id)


def create_unified_customer_preview(
    *,
    customer_code: str,
    factory_id: str,
    received_date: str,
    po_files: list[tuple[str, bytes]],
    schedule_file_name: str,
    schedule_content: bytes,
) -> dict[str, Any]:
    _validate_po_batch(po_files)
    ensure_unified_schedule(schedule_file_name, schedule_content, factory_id=factory_id, customer_code=customer_code)
    if not _factory_allowed(customer_code, factory_id):
        raise CustomerOrderUnifiedError(
            f"客户 {customer_code} 不属于当前厂区 {factory_id}"
        )
    try:
        normalized_received_date = date.fromisoformat(received_date).isoformat()
    except ValueError as exc:
        raise CustomerOrderUnifiedError("来单日期必须是 YYYY-MM-DD") from exc

    history = read_unified_history(schedule_content, factory_id=factory_id, customer_code=customer_code)
    if customer_code == 'ubtech':
        history = [h for h in history if '优必选' in h.customer_name or 'UBTECH' in h.customer_name.upper()]
    customer_name = _customer_name(customer_code, factory_id)
    try:
        rows, parser_warnings, input_template = _create_customer_rows(
            customer_code,
            factory_id,
            po_files,
            normalized_received_date,
            history,
        )
    except CustomerOrderUnifiedError:
        raise
    except CustomerOrderWorkbookError:
        raise
    except Exception as exc:
        raise CustomerOrderUnifiedError(str(exc)) from exc

    detail_rows = [row for row in rows if row.get("row_role") != "parent"]
    if factory_id == 'huakang-a' and customer_code == 'green-toys':
        workbook = _load_workbook(schedule_content)
        try:
            recovery = huakang_unified.green_summary_recovery(workbook, detail_rows, history)
        finally:
            workbook.close()
        assigned = {key: value['reference'] for key, value in recovery.items()}
        for row in detail_rows:
            record = {'contract_no': row['contract_no'], 'item_no': row['product_no']}
            huakang_unified.reconcile_green_po(record, history, assigned)
            row['reference_no'] = row['po_no'] = record['customer_po']
            key = (_text(row['contract_no']), _key(row['product_no']))
            if key in recovery:
                parser_warnings.append(f"{row['po_no']} / {row['product_no']}：两张摘要一致但缺少 ITEM 明细；导出时恢复明细和公式关联，保留原编号及人工价格。")
    huaxing_sheets = []
    if huaxing.enabled(factory_id, customer_code):
        workbook = _load_workbook(schedule_content)
        try:
            huaxing_sheets = huaxing.item_sheets(workbook, customer_code)
        finally:
            workbook.close()
    for row in detail_rows:
        row["received_date"] = normalized_received_date
        if huaxing_sheets:
            huaxing.prepare_row(row, history, customer_code, huaxing_sheets, factory_id)
        if regional.enabled(factory_id, customer_code):
            regional.reconcile_product(row, history, factory_id, customer_code)
        _enrich_row(row, history, customer_name)
        if huaxing.enabled(factory_id, customer_code):
            row['target_template'] = huaxing.target_template(factory_id)
        if huakang_unified.enabled(factory_id, customer_code):
            row["target_template"] = huakang_unified.TARGET_TEMPLATE
            for field in ("barcode", "port", "printing_requirement", "country"):
                if not _text(row.get(field)):
                    candidates = {item.extras[field] for item in history if _key(item.product_no) == _key(row.get("product_no")) and item.extras.get(field)}
                    if len(candidates) == 1:
                        row[field] = candidates.pop()
                        row.setdefault("lineage", {})[field] = f"ITEM表 · 货号 {row['product_no']} 唯一历史值"
            huakang_unified.round_cartons(row)
        elif regional.enabled(factory_id, customer_code):
            regional.enrich(row, history, factory_id, customer_code)
        if customer_code == 'ubtech':
            from app.services.customer_order_ubtech import finish_row
            finish_row(row)
    _mark_batch_duplicates(detail_rows)
    file_names = [name for name, _ in po_files]
    hashes = [sha256(content).hexdigest() for _, content in po_files]
    safe_stem = re.sub(r'[\\/:*?"<>|]+', "_", Path(schedule_file_name).stem).strip(" ._")
    output_file_name = f"{safe_stem or '河源业务统一排期'}_{customer_name}新单.xlsx"
    warnings = list(dict.fromkeys([
        "所有客户统一写入《河源业务统一排期》的 ITEM表、接单表和正单评审表。",
        "ITEM表 A:AB 为可自动写入范围；PO及可靠映射没有的可选字段保持空白。",
        "历史查重使用 SO#/Reference + 产品编号；源排期不会被覆盖。",
        *parser_warnings,
    ]))
    if huakang_unified.enabled(factory_id, customer_code):
        warnings[:3] = [
            "华康A三张表分别在各自取消单边界上方追加，同步 ITEM表 D:H 到接单表、正单评审表 C:G。",
            "通用字段及已明确映射的客户专属字段可自动写入；生产、出货、发票等人工字段保留。",
            "历史查重兼容未填SO的客户PO/Release编号；源排期不会被覆盖。",
        ]
    elif regional.enabled(factory_id, customer_code) or huaxing.enabled(factory_id, customer_code):
        warnings[:3] = [
            '各 ITEM 分类页独立追加，接单表和正单评审表关联实际明细页及行。',
            '只校验公共区域，客户专属字段按已确认的实际表头映射；人工生产和出货内容保留。',
            '历史查重兼容未填SO的客户PO编号，并读取取消单和已走货区中的历史订单。',
        ]
        if customer_code == 'edu':
            warnings[2] = 'EDU 按客户完整合同/PO＋产品查重，EDUHX 仅为内部编号；不同客户订单独立续编号，源排期不会被覆盖。'
    return {
        "preview_schema_version": PREVIEW_SCHEMA_VERSION,
        "customer_code": customer_code,
        "factory_id": factory_id,
        "po_file_name": file_names[0] if len(file_names) == 1 else f"{len(file_names)}个文件",
        "po_file_names": file_names,
        "po_file_count": len(file_names),
        "schedule_file_name": schedule_file_name,
        "source_po_sha256": _combined_hash(file_names, hashes),
        "source_po_sha256s": hashes,
        "source_schedule_sha256": sha256(schedule_content).hexdigest(),
        "input_template": input_template,
        "target_template": detail_rows[0]['target_template'] if customer_code == 'ubtech' else huakang_unified.TARGET_TEMPLATE if huakang_unified.enabled(factory_id, customer_code) else regional.TARGET_TEMPLATE if regional.enabled(factory_id, customer_code) else huaxing.target_template(factory_id) if huaxing.enabled(factory_id, customer_code) else TARGET_TEMPLATE,
        "output_file_name": output_file_name,
        "summary": {
            "total": len(detail_rows),
            "valid": sum(row["status"] == "valid" for row in detail_rows),
            "warning": sum(row["status"] == "warning" for row in detail_rows),
            "blocked": sum(row["status"] == "blocked" for row in detail_rows),
        },
        "rows": detail_rows,
        "warnings": warnings,
    }


def _copy_template_row(worksheet, source_row: int, target_row: int) -> None:
    worksheet.row_dimensions[target_row].height = worksheet.row_dimensions[source_row].height
    for column in range(1, worksheet.max_column + 1):
        source = worksheet.cell(source_row, column)
        target = worksheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.protection = copy(source.protection)
        target.alignment = copy(source.alignment)
        if isinstance(source.value, ArrayFormula):
            target.value = huakang_unified.copy_array_formula(source.value, source.coordinate, target.coordinate)
        elif isinstance(source.value, str) and source.value.startswith("="):
            try:
                target.value = Translator(
                    source.value,
                    origin=f"{get_column_letter(column)}{source_row}",
                ).translate_formula(f"{get_column_letter(column)}{target_row}")
            except Exception:
                target.value = source.value
        else:
            target.value = None
        coordinate = f"{get_column_letter(column)}{source_row}"
        for validation in worksheet.data_validations.dataValidation:
            if coordinate in validation.cells:
                validation.add(target.coordinate)


def _row_is_available(workbook, row: int) -> bool:
    item = workbook[ITEM_SHEET]
    order = workbook[ORDER_SHEET]
    review = workbook[REVIEW_SHEET]
    return (
        not _text(item.cell(row, 5).value)
        and not _text(item.cell(row, 8).value)
        and not _text(order.cell(row, 4).value)
        and not _text(order.cell(row, 7).value)
        and not _text(review.cell(row, 4).value)
        and not _text(review.cell(row, 7).value)
    )


def _ensure_output_slots(workbook, count: int) -> list[int]:
    markers = {name: _marker_row(workbook[name]) for name in SHEETS}
    if len(set(markers.values())) != 1:
        raise CustomerOrderUnifiedError("统一排期三张表的“取消单”边界行不一致")
    marker = next(iter(markers.values()))
    keyed_rows = [
        row for row in range(4, marker)
        if any((
            _text(workbook[ITEM_SHEET].cell(row, 5).value),
            _text(workbook[ITEM_SHEET].cell(row, 8).value),
            _text(workbook[ORDER_SHEET].cell(row, 4).value),
            _text(workbook[ORDER_SHEET].cell(row, 7).value),
            _text(workbook[REVIEW_SHEET].cell(row, 4).value),
            _text(workbook[REVIEW_SHEET].cell(row, 7).value),
        ))
    ]
    append_row = max(keyed_rows, default=3) + 1
    slots = [
        row for row in range(append_row, marker)
        if _row_is_available(workbook, row)
    ]
    missing = count - len(slots)
    for _ in range(max(0, missing)):
        for sheet_name in SHEETS:
            worksheet = workbook[sheet_name]
            current_marker = _marker_row(worksheet)
            insert_sparse_rows(worksheet, current_marker, 1)
            _copy_template_row(worksheet, current_marker - 1, current_marker)
        slots.append(_marker_row(workbook[ITEM_SHEET]) - 1)
    return slots[:count]


def _fit_generated_text(worksheet, row_number: int, linked_values: dict[int, Any] | None = None) -> None:
    """Expand only new rows when wrapped customer/product text needs more room."""
    height = worksheet.row_dimensions[row_number].height or 15
    for cell in worksheet[row_number]:
        value = (linked_values or {}).get(cell.column, cell.value)
        if not isinstance(value, str) or value.startswith('=') or not cell.alignment.wrap_text:
            continue
        dimension = worksheet.column_dimensions.get(get_column_letter(cell.column))
        if dimension is None:
            dimension = next((d for d in worksheet.column_dimensions.values() if d.min and d.max and d.min <= cell.column <= d.max), None)
        width = (dimension.width if dimension is not None else worksheet.sheet_format.defaultColWidth) or 13
        lines = sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(ch) in {'W', 'F'} else 1 for ch in line) / max(width - 2, 1))) for line in value.split('\n'))
        height = max(height, lines * (cell.font.sz or 11) * 1.35 + 4)
    worksheet.row_dimensions[row_number].height = min(height, 409)


def _write_item_row(worksheet, row_number: int, row: dict[str, Any]) -> None:
    values: dict[int, Any] = {
        1: _date_value(row.get("received_date")),
        2: _text(row.get("order_type")),
        3: _text(row.get("production_no")),
        4: _text(row.get("contract_no")),
        5: _text(row.get("reference_no")),
        6: _text(row.get("po_no")),
        7: _text(row.get("customer_name")) or _text(row.get("customer_country")),
        8: _text(row.get("product_no")),
        9: _text(row.get("product_name_zh")),
        10: _text(row.get("product_name_en")),
        11: _number(row.get("quantity")),
        12: _number(row.get("units_per_carton")),
        14: _text(row.get("standard")),
        15: _text(row.get("manual")),
        16: _text(row.get("label")),
        17: _text(row.get("customer_label")),
        18: _text(row.get("carton_mark")),
        19: _text(row.get("fabric_label")),
        20: _text(row.get("packaging")),
        21: _text(row.get("date_code")),
        22: _text(row.get("system_status")),
        23: _date_value(row.get("release_date")),
        24: _date_value(row.get("sample_date")),
        25: _date_value(row.get("customer_q") or row.get("line_q")),
        26: _date_value(row.get("requested_ship_date")),
        27: _text(row.get("third_party_inspection")),
        28: _text(row.get("inspection_result")),
    }
    for column, value in values.items():
        worksheet.cell(row_number, column).value = _excel_value(value) if value not in ("", None) else None
    quantity_cell = worksheet.cell(row_number, 11)
    units_cell = worksheet.cell(row_number, 12)
    if quantity_cell.value is not None and units_cell.value not in (None, 0):
        worksheet.cell(row_number, 13).value = (
            f"=CEILING(K{row_number}/L{row_number},1)" if row.get("_huakang_customer") or row.get('_regional_customer') == 'ubtech'
            else f"=K{row_number}/L{row_number}"
        )
    else:
        worksheet.cell(row_number, 13).value = None
    if row.get('_regional_customer') or row.get('_huaxing_customer'):
        for column in (11, 12, 13):
            worksheet.cell(row_number, column).number_format = 'General'
    if row.get('_huaxing_customer'):
        for column in (7, 9, 10, 14, 20):
            cell = worksheet.cell(row_number, column)
            alignment = copy(cell.alignment)
            alignment.wrapText = True
            cell.alignment = alignment
    if row.get('_regional_customer') or row.get('_huakang_customer') or row.get('_huaxing_customer'):
        _fit_generated_text(worksheet, row_number)


def _write_summary_row(worksheet, row_number: int, row: dict[str, Any], *, item_row: int | None = None, item_sheet=None) -> None:
    worksheet.cell(row_number, 3).value = _text(row.get("contract_no")) or None
    worksheet.cell(row_number, 4).value = _text(row.get("reference_no")) or None
    worksheet.cell(row_number, 5).value = _text(row.get("po_no")) or None
    worksheet.cell(row_number, 6).value = _text(row.get("customer_name")) or _text(row.get("customer_country")) or None
    worksheet.cell(row_number, 7).value = _text(row.get("product_no")) or None
    if item_row is not None:
        if row.get('_huaxing_customer'):
            for column in (6, 8, 9, 13):
                cell = worksheet.cell(row_number, column)
                alignment = copy(cell.alignment)
                alignment.wrapText = True
                cell.alignment = alignment
        # Generated rows have an exact ITEM counterpart. Key-based template
        # lookups intentionally go blank for duplicates and some supplied
        # reserved rows already contain #REF!, so bind only new summary rows.
        linked = {1: 1, 2: 2, 8: 9, 9: 10, 10: 11, 11: 12, 12: 13, 13: 14, 18: 25, 19: 26}
        if item_sheet is not None and _normalized_header(worksheet.cell(3, 20).value) == '备注':
            remarks = [c.column for c in item_sheet[3] if _normalized_header(c.value) == '备注']
            if len(remarks) == 1:
                linked[20] = remarks[0]
        for summary_column, item_column in linked.items():
            item_name = (item_sheet.title if item_sheet is not None else ITEM_SHEET).replace("'", "''")
            ref = f"'{item_name}'!{get_column_letter(item_column)}{item_row}"
            worksheet.cell(row_number, summary_column).value = f'=IF(LEN({ref})=0,"",{ref})'
        _fit_generated_text(worksheet, row_number, {c: item_sheet.cell(item_row, ic).value for c, ic in linked.items()} if item_sheet is not None else None)
        if row.get('_regional_customer') == 'goliath' and _number(row.get('unit_price_usd')) is not None:
            worksheet.cell(row_number, 16).value = _excel_value(_number(row['unit_price_usd']))
            worksheet.cell(row_number, 14).value = f'=P{row_number}*7.75'
            worksheet.cell(row_number, 15).value = f'=J{row_number}*N{row_number}'
            worksheet.cell(row_number, 17).value = f'=J{row_number}*P{row_number}'


def _recalculate_preview_row(row: dict[str, Any]) -> None:
    if not _text(row.get("reference_no")):
        row["reference_no"] = _text(row.get("contract_no")) or _text(row.get("po_no"))
    quantity = _number(row.get("quantity"))
    units = _number(row.get("units_per_carton"))
    price = _number(row.get("unit_price_hkd"))
    if quantity is not None and units not in (None, Decimal(0)):
        row["carton_count"] = _number_text(quantity / units)
    if quantity is not None and price is not None:
        row["amount_hkd"] = _number_text(quantity * price, 2)
    if row.get("_huakang_customer"):
        huakang_unified.round_cartons(row)
    if row.get('_regional_customer') == 'ubtech':
        from app.services.customer_order_ubtech import finish_row
        finish_row(row)


def _validate_resolutions(preview: dict[str, Any], requested: set[str]) -> None:
    available = {
        issue.get("skip_key")
        for row in preview["rows"]
        for issue in row["issues"]
        if issue.get("skip_key")
    }
    if requested - available:
        raise CustomerOrderUnifiedError("所选确认项已失效，请重新预览后再生成")
    blockers = [
        issue["message"]
        for row in preview["rows"]
        for issue in row["issues"]
        if issue.get("severity") == "blocked" and issue.get("skip_key") not in requested
    ]
    if blockers:
        raise CustomerOrderUnifiedError("仍有阻断项：" + "；".join(dict.fromkeys(blockers)))


def export_unified_customer_schedule(
    *,
    customer_code: str,
    factory_id: str,
    received_date: str,
    po_files: list[tuple[str, bytes]],
    schedule_file_name: str,
    schedule_content: bytes,
    skipped_issue_keys: set[str] | None = None,
    manual_overrides: list[dict[str, str]] | None = None,
) -> tuple[bytes, str, dict[str, Any]]:
    preview = create_unified_customer_preview(
        customer_code=customer_code,
        factory_id=factory_id,
        received_date=received_date,
        po_files=po_files,
        schedule_file_name=schedule_file_name,
        schedule_content=schedule_content,
    )
    decorate_manual_resolution_policy(preview)
    original_identities = {row['id']: tuple(_text(row.get(k)) for k in ('product_no', 'reference_no', 'contract_no', 'po_no')) for row in preview['rows']}
    apply_overrides_to_preview(preview, manual_overrides or [])
    for row in preview['rows']:
        _recalculate_preview_row(row)
    if huaxing.enabled(factory_id, customer_code) and manual_overrides:
        workbook = _load_workbook(schedule_content)
        try:
            huaxing.validate_edited_identities(preview, original_identities,
                huaxing.read_history(workbook, customer_code), huaxing.item_sheets(workbook, customer_code), manual_overrides or [])
        finally:
            workbook.close()
    for row in preview["rows"]:
        _recalculate_preview_row(row)
    requested = set(skipped_issue_keys or set())
    _validate_resolutions(preview, requested)

    workbook = _load_workbook(schedule_content)
    try:
        rows = [row for row in preview["rows"] if row.get("row_role") != "parent"]
        if huaxing.enabled(factory_id, customer_code):
            huaxing.export_rows(workbook, rows)
        elif huakang_unified.enabled(factory_id, customer_code) or regional.enabled(factory_id, customer_code):
            recovery = {}
            if factory_id == 'huakang-a' and customer_code == 'green-toys':
                recovery = huakang_unified.green_summary_recovery(workbook, rows,
                    read_unified_history(schedule_content, factory_id=factory_id, customer_code=customer_code))
            recovered_rows = [recovery.get((_text(row.get('contract_no')), _key(row.get('product_no')))) for row in rows]
            if any(recovered_rows):
                if any(value and (row['reference_no'] != value['reference'] or row['po_no'] != value['reference'])
                       for row, value in zip(rows, recovered_rows)):
                    raise CustomerOrderUnifiedError('恢复摘要时订单身份已改变，请重新核对原单')
                count = len(rows) - sum(bool(value) for value in recovered_rows)
                slots_by_sheet = huakang_unified.output_slots(workbook, len(rows), sheet_counts={
                    ITEM_SHEET: len(rows), ORDER_SHEET: count, REVIEW_SHEET: count,
                })
            else:
                slots_by_sheet = huakang_unified.output_slots(workbook, len(rows))
            summary_index = 0
            green_pairs = []
            for index, row in enumerate(rows):
                row_number = slots_by_sheet[ITEM_SHEET][index]
                _write_item_row(workbook[ITEM_SHEET], row_number, row)
                if huakang_unified.enabled(factory_id, customer_code):
                    huakang_unified.write_extras(workbook[ITEM_SHEET], row_number, row, customer_code)
                else:
                    regional.write_extras(workbook[ITEM_SHEET], row_number, row, factory_id, customer_code)
                if customer_code == 'ubtech':
                    from app.services.customer_order_ubtech import write_extensions
                    write_extensions(workbook[ITEM_SHEET], row_number, row)
                for name in (ORDER_SHEET, REVIEW_SHEET):
                    summary_row = recovered_rows[index][name] if recovered_rows[index] else slots_by_sheet[name][summary_index]
                    _write_summary_row(workbook[name], summary_row, row, item_row=row_number, item_sheet=workbook[ITEM_SHEET])
                    if factory_id == 'huakang-a' and customer_code == 'green-toys':
                        green_pairs.append((name, summary_row, row_number))
                    if customer_code == 'ubtech':
                        from app.services.customer_order_ubtech import write_summary
                        write_summary(workbook[name], slots_by_sheet[name][index], row, slots_by_sheet[ORDER_SHEET][index])
                if not recovered_rows[index]:
                    summary_index += 1
            for name, summary_row, item_row in green_pairs:
                detail = workbook[ITEM_SHEET]
                summary = workbook[name]
                if (not _text(detail.cell(item_row, 5).value) or not _text(detail.cell(item_row, 8).value)
                        or [_text(detail.cell(item_row, c).value) for c in range(4, 9)] !=
                           [_text(summary.cell(summary_row, c).value) for c in range(3, 8)]):
                    raise CustomerOrderUnifiedError('Green Toys 导出三表身份不一致，已停止生成，请重新核对')
                for c, ic in ((1,1),(2,2),(8,9),(9,10),(10,11),(11,12),(12,13),(13,14),(18,25),(19,26)):
                    ref = f"'ITEM表'!{get_column_letter(ic)}{item_row}"
                    if summary.cell(summary_row, c).value != f'=IF(LEN({ref})=0,"",{ref})':
                        raise CustomerOrderUnifiedError('Green Toys 导出三表公式关联不一致，已停止生成')
        else:
            slots = _ensure_output_slots(workbook, len(rows))
            for row_number, row in zip(slots, rows, strict=True):
                _write_item_row(workbook[ITEM_SHEET], row_number, row)
                _write_summary_row(workbook[ORDER_SHEET], row_number, row)
                _write_summary_row(workbook[REVIEW_SHEET], row_number, row)
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        output = BytesIO()
        workbook.save(output)
    finally:
        workbook.close()
    return output.getvalue(), preview["output_file_name"], preview
