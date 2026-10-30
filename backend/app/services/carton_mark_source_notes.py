"""Bounded source notes, separate from print fields and executable instructions."""
import io
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.services import carton_mark_ai_pilot as pilot
from app.services.carton_mark import extract_excel_document_items

MAX_REVIEW_SOURCE_NOTES = 2500
MAX_REVIEW_NOTES = 3000  # Native notes plus bounded formula/image/manual references.
MAX_REVIEW_NOTE_CHARACTERS = 120000


def collect(filename, content, extracted=None,worksheets=None,*,extended=False):
    extracted = extracted or extract_excel_document_items(filename, content)
    notes = []
    for cell in extracted.items:
        if re.fullmatch(r"[（(].*(?:印刷外箱|填写).*[)）]\s*(?:KGS|KG|LBS|CM)?", cell.text, re.I):
            continue  # An unfilled print field, not a separate production note.
        extended_requirement=extended and re.search(r"ZURU\s*LOGO|LOGO.*(?:中间|添加)|(?:add|print|place).*LOGO|条码.*(?:扫描|扫出来|不需.*扫)|barcode.*(?:scan|size)|^\s*This is an example\s*$",cell.text,re.I)
        if extended_requirement or re.search(r"备注|印刷|条码尺寸|按实际|LOGO.*(?:居中|放大)|REMARK|PRINTING\s+NOTE|^\s*\d+[.、．](?!\d)|字体|警告标识|纸箱要求|必须|不能低于|不小于|至少|需要加|需加|尽量.*居中|默认.*加|外箱.*(?:向上|朝上)|除了.*还需要", cell.text, re.I):
            notes.append((cell.location, cell.text))
    image_count = 0
    if Path(filename).suffix.lower() == ".xlsx":
        book = load_workbook(io.BytesIO(content), keep_links=False)
        try:
            for sheet in book:
                if sheet.sheet_state != "visible" or (worksheets and sheet.title not in worksheets):
                    continue
                image_count += len(sheet._images)
                for row in sheet:
                    for cell in row:
                        if (cell.comment and not sheet.row_dimensions[cell.row].hidden
                                and not sheet.column_dimensions[get_column_letter(cell.column)].hidden):
                            notes.append((f"{sheet.title}!{cell.coordinate}/comment", cell.comment.text))
        finally:
            book.close()
    unique = list(dict.fromkeys((location, pilot.normalized(text)) for location, text in notes if text.strip()))
    count_limit = MAX_REVIEW_SOURCE_NOTES if extended else 50
    character_limit = MAX_REVIEW_NOTE_CHARACTERS if extended else 16000
    if len(unique) > count_limit or any(len(text) > 4000 for _, text in unique) or sum(len(t) for _, t in unique) > character_limit:
        if extended:
            raise pilot.PilotError(f"所选工作表备注超过{count_limit}条、{character_limit}字或单条4000字，请缩小工作表范围；未调用API")
        raise pilot.PilotError("制版备注过多，请整理到单份原稿后重试")
    return [pilot.Evidence(id=f"N{i+1}", kind="current_note", location=location, text=text)
            for i, (location, text) in enumerate(unique)], image_count
