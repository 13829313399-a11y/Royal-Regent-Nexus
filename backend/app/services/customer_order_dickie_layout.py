"""Dickie exact-item/explicit-series placement; never consumes MO reservations."""
from collections import defaultdict
import re

from openpyxl.styles import PatternFill

# Explicit business-confirmed members only; never infer a family by truncation.
SERIES = {'203747029': '203747020', '203747020': '203747020'}
HISTORICAL_SECTIONS = {'取消单', '取消订单', '已走货', '已走货订单', '已完成', '已完成订单'}


def series_key(value):
    from app.services.customer_order_unified import _key
    product = _key(value)
    return SERIES.get(product, product)


def _columns(sheet):
    return (4, 7, 10) if sheet.title in ('接单表', '正单评审表') else (5, 8, 11)


def _stock(sheet, number):
    from app.services.customer_order_unified import _text
    from app.services.huaxing_order_legacy.edu_grouped_export import _green
    ref, product, quantity = _columns(sheet)
    kind = sheet.cell(number, 2)
    kind_text = _text(kind.value).upper() if kind.data_type != 'f' else ''
    reference = sheet.cell(number, ref)
    reference_text = _text(reference.value).upper() if reference.data_type != 'f' else ''
    return ('MO' in kind_text or '备料' in kind_text
            or reference_text in {'已入系统', '欠MO', '备料', '备料单'}
            or all(_green(sheet.cell(number, c)) for c in (ref, product, quantity)))


def active_boundary(sheet):
    """Normal new orders end before structural totals or historical sections."""
    from app.services.customer_order_unified import _text, CustomerOrderUnifiedError
    ref, product, _ = _columns(sheet)
    historical, totals = [], []
    summary = sheet.title in ('接单表', '正单评审表')
    for (row, column), cell in list(sheet._cells.items()):
        if row < 4 or cell.data_type == 'f' or not isinstance(cell.value, str):
            continue
        # A real order's fields/notes never become a section boundary merely
        # because they mention cancellation or shipment.
        if any(sheet.cell(row, c).data_type not in ('f', 'e')
               and sheet.cell(row, c).value not in (None, '') for c in (ref, product)):
            continue
        label = re.sub(r'\s+', '', _text(cell.value))
        if column <= 3 and label in HISTORICAL_SECTIONS:
            historical.append(row)
        elif summary and column == 10 and re.fullmatch(
                r'20\d{2}(?:[-–—/]20\d{2})?年接单(?:合计)?', label):
            totals.append(row)
    if not historical:
        raise CustomerOrderUnifiedError(f'{sheet.title} 未找到正常新单区与取消/已走货区的分界，请核对分区标题')
    return min(historical + totals)


def _placement(sheet, product_key):
    from app.services.customer_order_unified import _text
    _, product, _ = _columns(sheet)
    boundary = active_boundary(sheet)
    matches = [r for r in range(4, boundary)
               if sheet.cell(r, product).data_type != 'f'
               and series_key(sheet.cell(r, product).value) == product_key]
    stock = [r for r in matches if _stock(sheet, r)]
    if matches:
        start = min(stock) if stock else max(matches) + 1
    else:
        # A genuinely new item starts a new group at the active-area end.
        occupied = [r for r in range(4, boundary) if any(
            c.value not in (None, '') and c.data_type != 'f' for c in sheet[r])]
        start = max(occupied, default=3) + 1
    ordinary = [r for r in range(4, boundary)
                if _text(sheet.cell(r, product).value) and not _stock(sheet, r)
                and sheet.cell(r, product).data_type != 'f']
    same_item = [r for r in ordinary if series_key(sheet.cell(r, product).value) == product_key]
    choices = same_item or ordinary or matches or [max(4, start - 1)]
    source = min(choices, key=lambda r: abs(r - start))
    return start, sheet.cell(source, product)


def allocate_rows(workbook, rows):
    """Allocate each sheet independently; keep Cell handles across later inserts."""
    from app.services.customer_order_unified import _copy_template_row, CustomerOrderUnifiedError
    # Reuse the existing fully reference-aware insertion engine only. No EDU
    # numbering, matching, reservation selection or deduction code is called.
    from app.services.huaxing_order_legacy.edu_grouped_export import _insert, _green
    groups = defaultdict(list)
    for index, row in enumerate(rows):
        product = series_key(row.get('product_no'))
        if not product:
            raise CustomerOrderUnifiedError('Dickie 新单缺少完整货号，不能确定分组')
        groups[(row['item_sheet_name'], product)].append((index, row))
    handles = [dict() for _ in rows]
    for (lane, product), group in groups.items():
        for name in (lane, '接单表', '正单评审表'):
            sheet = workbook[name]
            start, source = _placement(sheet, product)
            _insert(workbook, sheet, start, len(group))
            ref_col, product_col, quantity_col = _columns(sheet)
            clear_stock_fill = all(_green(sheet.cell(source.row, c)) for c in (ref_col, product_col, quantity_col))
            for offset, (index, row) in enumerate(group):
                number = start + offset
                _copy_template_row(sheet, source.row, number)
                if clear_stock_fill:
                    for cell in sheet[number]:
                        if _green(cell):
                            cell.fill = PatternFill()
                # Seed keys now; a later group can safely inspect allocated rows.
                sheet.cell(number, ref_col).value = row.get('reference_no')
                sheet.cell(number, product_col).value = row.get('product_no')
                handles[index][name] = sheet.cell(number, product_col)
    return handles
