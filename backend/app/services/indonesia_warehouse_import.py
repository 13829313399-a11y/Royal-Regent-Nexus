"""Warehouse-document intake. Source facts, matching and posting stay separate."""
from collections import Counter
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import PurePath
from types import SimpleNamespace
import csv
import hashlib
import re
import unicodedata
from zipfile import ZipFile

import openpyxl
from fastapi import HTTPException
from pydantic import ValidationError
from app.schemas import indonesia_materials as s
from app.services.indonesia_material_calculations import ZERO, decimal_text, number, require
from app.services.indonesia_materials import Ledger, authorize, digest, now_text, view
from app.services.indonesia_material_options import automatic_billing
from app.services.identity_policy import lock_mutation
from app.services.transaction_lock import lock_transaction

TYPE = 'warehouse_receipt'
MAX_ROWS = 2000
FIELDS = {
    'factory': ['厂区', '厂区编码', '工厂'],
    'customer': ['客户', '客人'], 'recipient': ['接收方', '印尼接收方'],
    'warehouse': ['仓库', '收料仓库', '仓库名称'], 'receipt_no': ['收料单号', '入库单号', '实收单号', '送货单号'],
    'receipt_line': ['行号', '单据行号', '收料行号'], 'actual_date': ['实收日期', '收料日期', '入库日期'],
    'material_code': ['物料编码', '物料编号', '料号'], 'material_name': ['物料名称', '品名', '物料中文名称', '正式报关名称', '正式报关中文名称'],
    'specification': ['规格', '规格型号'], 'unit': ['单位', '数量单位', '计量单位'],
    'normal_quantity': ['正常实收数量', '正常数量', '正常量'], 'quantity': ['实收数量', '实收合计', '实际收到数量'],
    'spare_quantity': ['备品数量', '免费备品量', '备品量'], 'contract_no': ['合同号', '合同编号', '合同单号'],
    'item_no': ['货号', '成品货号', '产品编号'], 'batch': ['交付批次', '合同批次'],
    'po_no': ['客户PO', '客户PO号', '客户订单号'],
    'purchase_no': ['采购单号', '采购订单号'], 'purchase_line': ['采购行号', '采购单行号'],
    'purchase_date': ['采购日期', '下单日期'], 'purchase_quantity': ['采购总量', '采购数量'],
    'supplier': ['供应商', '供应商名称'], 'price': ['采购单价', '单价'], 'currency': ['币种', '采购币种'],
    'package_label': ['包装组', '箱号'], 'cartons': ['箱数', '纸箱数量'],
    'length_cm': ['长cm', '箱长cm', '长度cm'], 'width_cm': ['宽cm', '箱宽cm', '宽度cm'],
    'height_cm': ['高cm', '箱高cm', '高度cm'], 'carton_weight_kg': ['每箱纸箱重量kg', '每箱纸箱重kg'],
}
TEMPLATE = ['仓库', '收料单号', '行号', '实收日期', '物料编码', '物料名称', '规格', '单位', '正常实收数量', '备品数量',
            '合同号', '客户PO', '货号', '采购单号', '采购行号', '采购日期', '采购总量', '供应商', '采购单价', '币种',
            '包装组', '箱数', '长cm', '宽cm', '高cm', '每箱纸箱重量kg']


def normalized(value):
    return re.sub(r'[\s（）()_*＊]', '', unicodedata.normalize('NFKC', str(value or '')).strip()).casefold()


HEADERS = {normalized(alias): field for field, aliases in FIELDS.items() for alias in [field, *aliases]}


def value_key(value):
    return ' '.join(unicodedata.normalize('NFKC', str(value or '')).strip().split()).casefold()


def source_number(value, label):
    try:
        result = number(value)
    except Exception:
        require(False, label + '不是有效数字，请核对原单', 422)
    require(result.is_finite() and 0 <= result <= Decimal('1000000000000') and result.as_tuple().exponent >= -8,
            label + '超出有效范围或小数位数，请核对原单', 422)
    return result


def text(cell):
    value = cell.value if hasattr(cell, 'value') else cell
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    result = str(value).strip()
    if hasattr(cell, 'number_format') and isinstance(value, int) and re.fullmatch(r'0+', cell.number_format or ''):
        result = result.zfill(len(cell.number_format))
    return result


def parse_warehouse_receipts(content, filename):
    require(0 < len(content) <= 25 * 1024 * 1024, '实收单为空或超过 25 MB', 413)
    sheets, book = [], None
    try:
        if filename.lower().endswith('.csv'):
            decoded = None
            for encoding in ('utf-8-sig', 'gb18030'):
                try:
                    decoded = content.decode(encoding)
                    break
                except UnicodeError:
                    continue
            require(decoded is not None, 'CSV 编码无法识别，请另存为 UTF-8 CSV', 422)
            csv.field_size_limit(100000)
            rows = []
            for cells in csv.reader(StringIO(decoded)):
                require(len(rows) < MAX_ROWS + 30 and len(cells) <= 80, '实收单超过 2000 行或 80 列，请分批导入', 422)
                rows.append(cells)
            sheets.append(('实收单', rows))
        else:
            require(filename.lower().endswith('.xlsx'), '请上传 .xlsx 或 .csv 实收单', 422)
            with ZipFile(BytesIO(content)) as archive:
                require(sum(item.file_size for item in archive.infolist()) <= 150 * 1024 * 1024, '实收单解压后过大', 422)
            book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
            require(len(book.worksheets) <= 10, '实收单最多包含 10 个工作表，请拆分导入', 422)
            for sheet in book:
                require(sheet.max_row <= MAX_ROWS + 30 and sheet.max_column <= 80, '实收单超过 2000 行或 80 列，请分批导入', 422)
                sheets.append((sheet.title, list(sheet.iter_rows())))
        from app.services.indonesia_receipt_template import TITLE, REFERENCE_HEADERS, ENTRY_FIELDS
        result, empty_template = [], False
        for sheet, rows in sheets:
            header_index, fields = None, []
            for index, cells in enumerate(rows[:30]):
                mapped = [HEADERS.get(normalized(text(cell))) for cell in cells]
                if any(key in mapped for key in ('normal_quantity', 'quantity')) and any(key in mapped for key in ('material_code', 'material_name')):
                    header_index, fields = index, mapped
                    require(not any(count > 1 for field, count in Counter(key for key in fields if key).items()), f'{sheet} 关键表头重复，请核对原单', 422)
                    break
            if header_index is None:
                continue
            headers = {text(cell) for cell in rows[header_index]}
            order_template = REFERENCE_HEADERS <= headers and '正式报关中文名称' in headers and any(
                text(cells[0]) == TITLE for cells in rows[:header_index] if cells)
            for position, cells in enumerate(rows[header_index + 1:], header_index + 2):
                raw = [text(cell) for cell in cells]
                if not any(raw):
                    continue
                require(all(len(value) <= 1000 for value in raw), '原单单元格内容过长，请拆分或整理后导入', 422)
                source = {field: raw[i] for i, field in enumerate(fields) if field and i < len(raw)}
                if order_template and not any(source.get(field) for field in ENTRY_FIELDS):
                    empty_template = True
                    continue
                if normalized(source.get('material_name')) in {'合计', '总计'} and not source.get('receipt_no'):
                    continue
                issues = []
                if any(getattr(cell, 'data_type', '') in {'f', 'e'} or (isinstance(cell, str) and cell.startswith('=')) for i, cell in enumerate(cells) if i < len(fields) and fields[i]):
                    issues.append('来源含公式或错误值，请在原单粘贴为数值后导入')
                for field in ('receipt_no', 'receipt_line', 'actual_date'):
                    if not source.get(field):
                        issues.append({'receipt_no': '缺少收料单号', 'receipt_line': '缺少稳定单据行号', 'actual_date': '缺少实收日期'}[field])
                if not source.get('material_code') and not source.get('material_name'):
                    issues.append('缺少物料编码和名称')
                try:
                    source['actual_date'] = date.fromisoformat(source.get('actual_date', '').replace('/', '-')).isoformat()
                except ValueError:
                    issues.append('实收日期须为明确的年月日，例如 2026-10-22')
                try:
                    spare = number(source.get('spare_quantity') or '0')
                    normal = number(source['normal_quantity']) if source.get('normal_quantity') else number(source['quantity']) - spare
                    require(normal.is_finite() and spare.is_finite() and normal >= 0 and spare >= 0 and normal + spare > 0, '数量无效', 422)
                    require(max(normal, spare) <= Decimal('1000000000000') and min(normal.as_tuple().exponent, spare.as_tuple().exponent) >= -8, '数量超限', 422)
                    require(not source.get('quantity') or number(source['quantity']) == normal + spare, '实收合计须等于正常量加备品量', 422)
                    source.update(normal_quantity=decimal_text(normal), spare_quantity=decimal_text(spare))
                except Exception:
                    issues.append('实收数量须为非负数、最多八位小数且合计大于零；实收合计包含备品')
                for key in ('price', 'purchase_quantity'):
                    if source.get(key):
                        try:
                            value = source_number(source[key], '采购单价' if key == 'price' else '采购总量')
                            require(key != 'purchase_quantity' or value > 0, '采购总量必须大于零', 422)
                            source[key] = decimal_text(value)
                        except HTTPException as error:
                            issues.append(str(error.detail))
                require(len(result) < MAX_ROWS, '单次最多导入 2000 行实收记录', 422)
                require(all(len(value) <= 500 for value in source.values()), '实收字段过长，请整理原表', 422)
                result.append(dict(index=len(result), source_sheet=sheet, source_row=position, source=source,
                                   source_issues=issues, raw_headers=[text(c) for c in rows[header_index]], raw_values=raw))
        require(result, '模板尚未填写本次实收数量，请填写后再导入。' if empty_template else
                '未找到实收表头；请使用导入模板或包含物料、实收数量的常见表头', 422)
        return result
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(422, '无法读取实收单，请核对文件格式及字段') from error
    finally:
        if book:
            book.close()


def identity(warehouse, source):
    line = value_key(source.get('receipt_line'))
    return digest([TYPE, value_key(source.get('warehouse') or warehouse), value_key(source.get('receipt_no')), str(int(line)) if line.isdigit() else line])


def candidates(ledger, record, source, material):
    result = []
    for order in ledger.rows('order'):
        if order.data['customer'] != record.data['customer'] or order.data['recipient'] != record.data['recipient']:
            continue
        if any(source.get(key) and source[key] != order.data.get(key, '') for key in ('contract_no', 'po_no', 'item_no', 'batch')):
            continue
        for demand in order.data['demands']:
            if demand['material_id'] != material.id or value_key(demand['unit']) != value_key(material.data['unit']):
                continue
            assigned = sum((number(row.data['normal_quantity']) + number(row.data['spare_quantity']) for row in ledger.rows('allocation') if row.data['order_id'] == order.id and row.data['demand_id'] == demand['id']), ZERO)
            remaining = max(number(demand['quantity']) - assigned, ZERO)
            result.append(dict(order_id=order.id, demand_id=demand['id'], order_version=order.version, remaining=decimal_text(remaining),
                               label=f"PO {order.data.get('po_no') or '未填写'} · 合同 {order.data['contract_no']} / {order.data['item_no']} · 走货 {order.data['delivery_date']}" +
                                     (f"（历史批次 {order.data['batch']}）" if order.data.get('batch') else '') +
                                     f" · 待分配 {decimal_text(remaining)} {demand['unit']}"))
            if len(result) > 200:
                return result  # The caller marks this row for review without failing the factory workspace.
    return result


def reconcile(ledger, record):
    rows = deepcopy(record.data['rows'])
    materials = ledger.rows('material')
    seen = Counter(identity(record.data['warehouse'], row['source']) for row in rows)
    posted = {row.data['warehouse_source']['identity'] for row in ledger.rows('supply') if row.data.get('warehouse_source')}
    candidate_count = 0
    for row in rows:
        source = row['source']
        row.update(issues=list(row['source_issues']), materials=[], targets={}, supplies=[])
        if candidate_count > 20000 or row.get('imported_supply_id'):
            if not row.get('imported_supply_id'):
                row['issues'].append('本批候选过多，请分批或补齐合同、货号、物料资料后核对')
            row.update(blocked=bool(row['issues']), auto_material_id=None)
            continue
        factory_aliases = {'huaxing': ['华兴'], 'huadeng': ['华登'], 'huakang-a': ['华康A', '华康 A', 'A厂'],
                           'huakang-b': ['华康B', '华康 B', 'B厂'], 'huakang-c': ['华康C', '华康 C', 'C厂'], 'huakang-d': ['华康D', '华康 D', 'D厂']}
        if source.get('factory') and value_key(source['factory']) not in {value_key(item) for item in [ledger.factory, *factory_aliases[ledger.factory]]}:
            row['issues'].append('原单厂区与当前厂区不一致或无法识别，请核对原单')
        if source.get('customer') and value_key(source['customer']) != value_key(record.data['customer']):
            row['issues'].append('原单客户与本次实收客户不一致，请核对上传选项')
        if source.get('recipient') and value_key(source['recipient']) != value_key(record.data['recipient']):
            row['issues'].append('原单接收方与本次实收接收方不一致，请核对上传选项')
        if not (source.get('warehouse') or record.data['warehouse']).strip():
            row['issues'].append('缺少仓库名称')
        if seen[identity(record.data['warehouse'], source)] > 1:
            row['issues'].append('同一仓库、收料单号和行号重复，请核对原单')
        if identity(record.data['warehouse'], source) in posted and not row.get('imported_supply_id'):
            row['issues'].append('该仓库收料单行已经导入，不能重复计入实收')
        found = [item for item in materials if (value_key(item.data['name']) == value_key(source['material_name']) if source.get('material_name') else value_key(item.data['code']) == value_key(source.get('material_code')))
                 and (not source.get('unit') or value_key(item.data['unit']) == value_key(source['unit']))
                 and (not source.get('specification') or value_key(item.data['specification']) == value_key(source['specification']))]
        if len(found) > 1 and source.get('material_code'):
            by_code = [item for item in found if value_key(item.data['code']) == value_key(source['material_code'])]
            if by_code:
                found = by_code
        if not found and source.get('material_name') and source.get('material_code'):
            # Keep legacy code/name discrepancies reviewable, never automatically accepted.
            found = [item for item in materials if value_key(item.data['code']) == value_key(source['material_code'])
                     and (not source.get('unit') or value_key(item.data['unit']) == value_key(source['unit']))
                     and (not source.get('specification') or value_key(item.data['specification']) == value_key(source['specification']))]
        if len(found) > 100:
            row['issues'].append('同名物料候选过多，请补齐编码、规格或单位后核对')
            found = []
        for material in found:
            row['materials'].append(dict(id=material.id, version=material.version, label=f"{material.data['code']} · {material.data['name']} · {material.data['specification']} · {material.data['unit']}",
                                         defaults={key: material.data[key] for key in ('supplier', 'price', 'currency')}))
            row['targets'][material.id] = candidates(ledger, record, source, material)
            if len(row['targets'][material.id]) > 200:
                row['targets'][material.id] = []
                row['issues'].append('合同候选过多，请在原单补充合同、货号或交付批次后核对')
            row['supplies'] += [dict(id=item.id, kind='supply', version=item.version, data={key: item.data[key] for key in
                                ('material_id', 'purchase_no', 'purchase_line', 'purchase_date', 'purchase_quantity', 'normal_quantity', 'spare_quantity', 'supplier', 'price', 'currency', 'batch')})
                                for item in ledger.rows('supply') if item.data['material_id'] == material.id and item.data['status'] == 'expected'
                                and (not source.get('purchase_no') or item.data['purchase_no'] == source['purchase_no'])]
            candidate_count += len(row['targets'][material.id])
        candidate_count += len(found) + len(row['supplies'])
        if candidate_count > 20000:
            row.update(materials=[], targets={}, supplies=[])
            row['issues'].append('本批候选过多，请分批或补齐合同、货号、物料资料后核对')
            found = []
        exact = len(found) == 1 and (bool(source.get('material_code')) or (value_key(found[0].data['specification']) == value_key(source.get('specification')) and bool(source.get('unit'))))
        if exact and source.get('material_name') and value_key(source['material_name']) != value_key(found[0].data['name']):
            exact = False  # A code/name disagreement requires an explicit business choice.
        if exact and source.get('material_code') and value_key(source['material_code']) != value_key(found[0].data['code']):
            exact = False
        row['auto_material_id'] = found[0].id if exact else None
        row['blocked'] = bool(row['issues'])
    return rows


def preview(db, user, factory, customer, recipient, warehouse, filename, content, rows, refresh_user):
    ledger = Ledger(db, factory)
    try:
        lock_mutation(db)
        lock_transaction(db, 'indonesia-shared-options', 'shared-options')
        lock_transaction(db, 'indonesia-materials', factory)
        current = refresh_user()
        require(current.id == user.id, '账号身份已变化，请重新核对', 401)
        authorize(current, factory, 'write')
        ledger.fixed_customer(customer)
        key = [TYPE, hashlib.sha256(content).hexdigest(), customer, recipient, warehouse]
        existing = next((item for item in ledger.rows('import') if item.business_key == digest(key)), None)
        if not existing:
            asset = ledger.asset(PurePath(filename).name, content, 'text/csv' if filename.lower().endswith('.csv') else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            existing = ledger.save('import', dict(import_type=TYPE, status='preview', customer=customer, recipient=recipient, warehouse=warehouse,
                filename=asset.filename, asset_id=asset.id, provided_by=current.id, rows=rows, shipment_ids=[]), key=key)
        result = view(existing)
        result['data']['rows'] = reconcile(ledger, existing)
        db.commit()
        return dict(record=result)
    except Exception:
        db.rollback()
        raise


def confirm(ledger, record, payload, actor):
    require(record.data.get('import_type') == TYPE, '请选择仓库实收导入记录', 422)
    selected = s.WarehouseSelection.model_validate(payload)
    require(len({row.index for row in selected.rows}) == len(selected.rows), '不能重复选择同一来源行', 422)
    ledger.fixed_customer(record.data['customer'])
    sources = {row['index']: row for row in reconcile(ledger, record)}
    packing_occurrences = Counter((value_key(item['source'].get('warehouse') or record.data['warehouse']), value_key(item['source'].get('receipt_no')),
                                   value_key(item['source'].get('package_label'))) for item in sources.values() if item['source'].get('package_label'))
    shipment_lines, packages, all_allocations = [], [], set()
    for choice in selected.rows:
        source_row = sources.get(choice.index)
        require(source_row and not source_row['blocked'] and not source_row.get('imported_supply_id'), '来源行有问题或已导入，请刷新核对', 409)
        source = source_row['source']
        material = ledger.get(choice.material_id, 'material')
        require(material.version == choice.material_version, '物料资料已更新，请刷新匹配后确认', 409)
        require(material.id in {item['id'] for item in source_row['materials']}, '所选物料与原单编码、名称、规格或单位不符', 422)
        normal, spare = number(source['normal_quantity']), number(source['spare_quantity'])
        require(sum((item.normal_quantity for item in choice.allocations), ZERO) == normal and sum((item.spare_quantity for item in choice.allocations), ZERO) == spare,
                '分配的正常量和备品量须分别等于该行实际收到数量', 422)
        require(len({(item.order_id, item.demand_id) for item in choice.allocations}) == len(choice.allocations), '同一来源行不能重复选择合同需求', 422)
        purchase_fields = {'purchase_no', 'purchase_line', 'purchase_date', 'purchase_quantity', 'supplier', 'price', 'currency'}
        require(set(choice.purchase) <= purchase_fields, '采购补录包含不支持的字段', 422)
        if choice.supply_id:
            supply = ledger.get(choice.supply_id, 'supply')
            require(supply.version == choice.supply_version and supply.data['status'] == 'expected' and supply.data['material_id'] == material.id, '原采购批次已变化或不适用于此物料', 409)
            require(not choice.purchase, '选用原采购批次时保留其原计价资料', 422)
            for key in ('purchase_no', 'purchase_line', 'supplier', 'currency'):
                require(not source.get(key) or value_key(source[key]) == value_key(supply.data[key]), '所选采购批次与原单采购或供应商资料不符', 422)
            for key in ('price', 'purchase_quantity'):
                require(not source.get(key) or source_number(source[key], '原单采购数字') == number(supply.data[key]), '原单单价或采购总量与原采购批次不一致，请核对；不会覆盖历史采购', 422)
            require(not source.get('purchase_date') or source['purchase_date'] == supply.data['purchase_date'], '原单采购日期与原采购批次不一致，请核对', 422)
            require(number(supply.data['normal_quantity']) == normal and number(supply.data['spare_quantity']) == spare,
                    '实际数量与原采购批次不一致，请先核对原批次数量或拆分回料批次', 422)
        else:
            purchase = {key: value for key, value in source.items() if key in purchase_fields and value != ''}
            purchase = {**{key: material.data[key] for key in ('supplier', 'price', 'currency')}, **purchase, **choice.purchase}
            # These are real purchase facts: missing identifiers/date/total must be supplied, never invented from receipt numbers.
            data = {**purchase, 'material_id': material.id, 'normal_quantity': str(normal), 'spare_quantity': str(spare),
                    'batch': 'WH-' + identity(record.data['warehouse'], source)[:24], 'billing_mode': 'automatic', 'billing_weight_source': 'net_weight_kg'}
            existing_purchase = [item for item in ledger.rows('supply') if all(item.data[key] == data.get(key) for key in ('supplier', 'purchase_no', 'purchase_line'))]
            require(not any(item.data['status'] == 'expected' for item in existing_purchase), '该采购单行已有待回料批次，请选择原批次，避免重复登记', 409)
            supply = ledger.execute(SimpleNamespace(action='supply_save', entity_id=None, expected_version=0, data=data), actor)
        evidence = dict(identity=identity(record.data['warehouse'], source), warehouse=source.get('warehouse') or record.data['warehouse'],
                        receipt_no=source['receipt_no'], receipt_line=source['receipt_line'], import_id=record.id,
                        asset_id=record.data['asset_id'], sheet=source_row['source_sheet'], row=source_row['source_row'])
        require(not any(item.data.get('warehouse_source', {}).get('identity') == evidence['identity'] for item in ledger.rows('supply')), '该仓库实收单行已导入', 409)
        supply_data = deepcopy(supply.data)
        supply_data.update(status='received', actual_date=source['actual_date'], confirmed_by=actor, confirmed_at=now_text(),
                           receipt_confirmation_source='sales_warehouse_document', warehouse_source=evidence)
        ledger.save('supply', supply_data, supply)
        if choice.replace_allocations:
            require(choice.supply_id and choice.replacement_reason.strip(), '重新指定历史分配须选择原采购批次并填写原因', 422)
            old = [item for item in ledger.rows('allocation') if item.data['supply_id'] == supply.id and
                   number(item.data['normal_quantity']) + number(item.data['spare_quantity']) > 0]
            refs = {item.allocation_id: item.version for item in choice.replace_allocations}
            require(len(refs) == len(choice.replace_allocations) and refs == {item.id: item.version for item in old},
                    '原采购批次的历史分配已变化，请重新核对全部旧分配', 409)
            require(not any(line['allocation_id'] in refs for shipment in ledger.rows('shipment') for line in shipment.data['lines']),
                    '旧分配已用于出货资料，请先核对原出货单，不能在实收导入中替换', 409)
            for item in old:
                changed = deepcopy(item.data)
                changed.update(normal_quantity='0', spare_quantity='0', replacement_reason=choice.replacement_reason,
                               replaced_by_import_id=record.id, replaced_by=actor, replaced_at=now_text())
                ledger.save('allocation', changed, item)
        row_lines = []
        for plan in choice.allocations:
            order, demand = ledger.demand(plan.order_id, plan.demand_id)
            require(order.version == plan.order_version, '合同需求已更新，请重新核对匹配', 409)
            require(order.data['customer'] == record.data['customer'] and order.data['recipient'] == record.data['recipient'], '只能选择本次客户及接收方的合同', 422)
            require(not any(source.get(key) and source[key] != order.data.get(key, '') for key in ('contract_no', 'po_no', 'item_no', 'batch')), '所选客户 PO、合同或货号与原单不符', 422)
            prior = [item for item in ledger.rows('allocation') if item.data['supply_id'] == supply.id and item.data['order_id'] == order.id and item.data['demand_id'] == demand['id']
                     and number(item.data['normal_quantity']) + number(item.data['spare_quantity']) > 0]
            if prior:
                require(len(prior) == 1 and all(number(prior[0].data[field]) == number(getattr(plan, field)) for field in ('normal_quantity', 'spare_quantity')), '原批次已有分配，请沿用原分配数量或先核对调整', 409)
                allocation = prior[0]
            else:
                allocation = ledger.execute(SimpleNamespace(action='allocate', entity_id=None, expected_version=0,
                    data=dict(supply_id=supply.id, order_id=order.id, demand_id=demand['id'], normal_quantity=str(plan.normal_quantity), spare_quantity=str(plan.spare_quantity), reason='仓库实收导入确认')), actor)
            require(allocation.id not in all_allocations, '同一分配行不可重复加入本次出货草稿', 422)
            all_allocations.add(allocation.id)
            line = dict(allocation_id=allocation.id, normal_quantity=str(plan.normal_quantity), spare_quantity=str(plan.spare_quantity))
            shipment_lines.append(line)
            row_lines.append(line)
        source_row['imported_supply_id'] = supply.id
        # Only complete, single-receipt packing is imported automatically. Mixed boxes are completed explicitly in the shipment editor.
        packing_keys = ('package_label', 'cartons', 'length_cm', 'width_cm', 'height_cm', 'carton_weight_kg')
        packing_identity = (value_key(source.get('warehouse') or record.data['warehouse']), value_key(source.get('receipt_no')), value_key(source.get('package_label')))
        if packing_occurrences[packing_identity] == 1 and all(source.get(key) != '' and source.get(key) is not None for key in packing_keys):
            try:
                package = s.Package.model_validate({**{key: source[key] for key in packing_keys if key != 'package_label'}, 'label': source['package_label'], 'items': row_lines}).model_dump(mode='json')
                packages.append(package)
            except ValidationError:
                source_row.setdefault('packing_notes', []).append('原单装箱资料无效，请在出货草稿补填')
    # Repeated labels may describe a mixed carton, never infer multiple physical boxes from them.
    counts = Counter(item['label'] for item in packages)
    packages = [item for item in packages if counts[item['label']] == 1]
    shipment = ledger.execute(SimpleNamespace(action='shipment_save', entity_id=None, expected_version=0, data=dict(customer=record.data['customer'],
        recipient=record.data['recipient'], shipment_date=selected.shipment_date.isoformat(), lines=shipment_lines, packages=packages,
        remarks='由仓库实收单生成，请补齐并核对装箱资料后确认出货。')), actor)
    shipment_data = deepcopy(shipment.data)
    shipment_data.update(warehouse_import_id=record.id, packing_pending=len(packages) < len(selected.rows))
    ledger.save('shipment', shipment_data, shipment)
    data = deepcopy(record.data)
    data.update(rows=[sources[index] for index in sorted(sources)], status='partial' if any(not item.get('imported_supply_id') for item in sources.values()) else 'confirmed',
                shipment_ids=[*data.get('shipment_ids', []), shipment.id], confirmed_by=actor)
    return ledger.save('import', data, record)
