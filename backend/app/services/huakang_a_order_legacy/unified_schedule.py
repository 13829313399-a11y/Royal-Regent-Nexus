"""Huakang A's customer extensions to the September unified schedule."""
from __future__ import annotations

from app.services.sparse_worksheet import insert_rows as insert_sparse_rows

from copy import copy
from decimal import Decimal, ROUND_CEILING
import re

from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.formula import ArrayFormula


TARGET_TEMPLATE = "HEYUAN_BUSINESS_UNIFIED_HUAKANG_A_V4"

ITEM_TAILS = {
    "360": ("条码", "MS Container Type (FCL/LCL)", "Consolidated ID", "Port of Discharge", "备注", "版本"),
    "green-toys": (
        "走货国家", "落货港", "完成状态", "走货方式", "单价US", "货价总金额US$",
        "PO单价HK$", "货价总金额HK$", "装好纸箱报库存价HK$", "出厂总金额HK$",
        "走货时间", "柜号", "走货方式", "发票", "打票单价", "打票总金额", "备注",
    ),
    "headstart": (
        "条码", "印字要求", "走货国家", "订单状态（是否齐物料）", "订单单价USD",
        "总金额USD", "单价HK$", "总金额HK$", "装好纸箱报库存价HK$", "走货日期",
        "走货方式", "车头注明", "PDQ差价", "发票单价", "发票金额", "打票日期",
        "发票号码", "CBM", "请款金额", "装箱清单", "备注",
    ),
}
EXTRA_HEADERS = {
    "360": {
        "barcode": "条码", "container_type": "MS Container Type (FCL/LCL)",
        "consolidated_id": "Consolidated ID", "port": "Port of Discharge",
        "product_version": "版本",
        "shipping_item": "出货货号",
    },
    "green-toys": {
        "country": "走货国家", "port": "落货港", "transportation_mode": "走货方式",
        "unit_price_usd": "单价US", "amount_usd": "货价总金额US$",
        "unit_price_hkd": "PO单价HK$", "amount_hkd": "货价总金额HK$",
    },
    "headstart": {
        "barcode": "条码", "printing_requirement": "印字要求", "country": "走货国家",
        "unit_price_usd": "订单单价USD", "amount_usd": "总金额USD",
        "unit_price_hkd": "单价HK$", "amount_hkd": "总金额HK$",
    },
}


def extra_columns(sheet, customer_code):
    """Resolve only known, unambiguous customer fields outside the common area."""
    from app.services.customer_order_unified import _normalized_header
    found = {}
    for field, label in EXTRA_HEADERS.get(customer_code, {}).items():
        found[field] = [c.column for c in sheet[3] if c.column > 28
                        and _normalized_header(c.value) == _normalized_header(label)]
    # Green Toys has a planned mode before prices and an actual shipping mode
    # later. Only use that established distinction when it remains unambiguous.
    if customer_code == "green-toys":
        prices = found.get("unit_price_usd", [])
        ports = found.get("port", [])
        if len(prices) == len(ports) == 1:
            found["transportation_mode"] = [c for c in found["transportation_mode"] if ports[0] < c < prices[0]]
        else:
            found["transportation_mode"] = []
    return {field: columns[0] for field, columns in found.items() if len(columns) == 1}


def enabled(factory_id, customer_code):
    return factory_id == "huakang-a" and customer_code in ITEM_TAILS


def organize_360_schedule(workbook, *, source_content=None, today=None):
    """Only the Huakang A / 360 caller invokes this layout-only operation."""
    from .inspection_archive import organize
    return organize(workbook, source_content=source_content, today=today)


def history_reference(customer_code, contract, reference, po):
    if customer_code in {"360", "green-toys"}:
        return reference or po or contract
    return reference or contract or po


def identity(customer_code, reference, product):
    from app.services.customer_order_unified import _key
    if customer_code == "green-toys":
        reference = re.sub(r"-\d+$", "", str(reference or ""))
    elif customer_code == "headstart":
        from app.services.huakang_a_order_legacy.green_toys_headstart import _headstart_contract_key
        reference = _headstart_contract_key(reference)
    return _key(reference), _key(product)


def reconcile_green_po(record, history, assigned):
    from app.services.customer_order_unified import _key
    base = str(record.get("contract_no") or "")
    key = (base, _key(record.get("item_no")))
    matching = {h.po_no for h in history if h.po_no and identity("green-toys", h.po_no, h.product_no) == identity("green-toys", base, record.get("item_no"))}
    if len(matching) == 1:
        assigned[key] = next(iter(matching))
    if key not in assigned:
        used = [h.po_no for h in history] + list(assigned.values())
        suffixes = [int(m.group(1)) for value in used if (m := re.fullmatch(re.escape(base) + r"-(\d+)", value))]
        assigned[key] = f"{base}-{max(suffixes, default=0) + 1}" if base else ""
    record["customer_po"] = assigned[key]


def green_summary_recovery(workbook, rows, history):
    """Recover only two agreeing summary rows with no corresponding ITEM order.

    They reserve their existing PO suffix, not an order/quantity history record.
    Literal user data in derived fields makes recovery ambiguous and is refused.
    """
    from app.services.customer_order_unified import _key, _text, _marker_row, CustomerOrderUnifiedError
    keys = {(_text(r.get('contract_no')), _key(r.get('product_no'))) for r in rows}
    bases = {key[0] for key in keys}
    actual = {identity('green-toys', h.po_no or h.reference_no, h.product_no) for h in history}
    existing_references = {}
    for h in history:
        existing_references.setdefault(h.po_no or h.reference_no, set()).add(_key(h.product_no))
    found = {name: {} for name in ('接单表', '正单评审表')}
    for name in found:
        sheet = workbook[name]
        for n in range(4, _marker_row(sheet)):
            values = tuple(_text(sheet.cell(n, c).value) for c in range(3, 8))
            base, reference, po, customer, product = values
            if base not in bases:
                continue
            key = (base, _key(product))
            if identity('green-toys', base, product) in actual:
                continue
            if key not in keys:
                raise CustomerOrderUnifiedError(f'{name} 第{n}行：摘要缺少 ITEM 明细且不在当前原单中，请人工核对')
            if (key in found[name] or reference != po or not re.fullmatch(re.escape(base) + r'-\d+', po)
                    or _key(customer) not in {'GT', 'GREENTOYS'}):
                raise CustomerOrderUnifiedError(f'{name} 第{n}行：摘要身份不唯一，不能自动恢复')
            if existing_references.get(po, set()) - {key[1]}:
                raise CustomerOrderUnifiedError(f'{name} 第{n}行：摘要编号已被另一货号占用，不能自动恢复')
            for c in (1, 2, 8, 9, 10, 11, 12, 13, 18, 19, 20):
                cell = sheet.cell(n, c)
                if cell.value not in (None, '') and cell.data_type != 'f':
                    raise CustomerOrderUnifiedError(f'{name} 第{n}行：摘要有人工填写的派生字段，不能自动覆盖')
            found[name][key] = (n, values)
    if found['接单表'].keys() != found['正单评审表'].keys():
        raise CustomerOrderUnifiedError('Green Toys 两张摘要表的孤立订单不一致，请人工核对')
    result = {}
    references = set()
    for key, (n, values) in found['接单表'].items():
        review_n, review_values = found['正单评审表'][key]
        if values != review_values or values[2] in references:
            raise CustomerOrderUnifiedError('Green Toys 两张摘要表的订单编号/货号不一致，请人工核对')
        references.add(values[2])
        result[key] = {'reference': values[2], '接单表': n, '正单评审表': review_n}
    return result


def map_record(row, record, customer_code):
    """Separate actual business fields from the legacy 17-column display strings."""
    row["standard"] = record.get("standard") or ""
    row["packaging"] = record.get("product_packaging") or ""
    row["unit_price_usd"] = record.get("unit_price_usd")
    row["_huakang_customer"] = customer_code
    for field in EXTRA_HEADERS[customer_code]:
        if field not in {"amount_usd", "unit_price_hkd", "amount_hkd"}:
            row[field] = record.get(field) if record.get(field) is not None else ""
    if customer_code == "360":
        # RL Order Type is the customer's routing vocabulary (e.g. DI Order),
        # not the schedule's business category. New release orders are 正单;
        # any explicit business correction is applied after this source map.
        source_order_type = record.get("order_type") or ""
        row["order_type"] = "正单"
        # Customer Release (including OH numbers), the complete RL and the
        # customer's PO are separate identifiers, even when a value is absent.
        row["production_no"] = record.get("customer_release") or ""
        row["contract_no"] = record.get("contract_no") or ""
        row["reference_no"] = record.get("contract_no") or ""
        row["po_no"] = record.get("customer_po") or ""
        row["customer_release_no"] = record.get("customer_release") or ""
        row["customer_name"] = record.get("customer_name") or "ThreeSixty"
        source = row.get("source_po_file_name") or record.get("file_name") or "PO"
        row.setdefault("lineage", {}).update({
            "order_type": "华康A 360 业务规则 · 新 PURCHASE ORDER RELEASE 记为正单，客户 RL Order Type 不作为排期订单类型",
            "source_order_type": f"{source} · RL Order Type: {source_order_type or '未提供'}",
            "production_no": f"{source} · Customer Release No.（客户 Release/OH 号）",
            "contract_no": f"{source} · PURCHASE ORDER RELEASE 标题的完整 RL 编号",
            "reference_no": f"{source} · PURCHASE ORDER RELEASE 标题的完整 RL 编号",
            "po_no": f"{source} · Customer PO Number",
            "customer_release_no": f"{source} · Customer Release No.（保留来源值）",
            "customer_po_number": f"{source} · Customer PO Number: {record.get('customer_po') or ''}",
        })
        sticker_sources = []
        for evidence in record.get("outer_carton_sticker_sources") or []:
            if not isinstance(evidence, dict) or not evidence.get("raw_text"):
                continue
            parts = [str(evidence.get("file_name") or evidence.get("relative_path") or source)]
            if evidence.get("page") is not None:
                parts.append(f"第{evidence['page']}页")
            parts.append(str(evidence["raw_text"]))
            sticker_sources.append(" · ".join(parts))
        if sticker_sources:
            row["lineage"]["customer_label"] = "\n".join(sticker_sources)
    elif customer_code == "green-toys":
        row["reference_no"] = record.get("customer_po") or record.get("contract_no") or ""
        row["transportation_mode"] = "40'YT"
    else:
        row["po_no"] = record.get("contract_no") or ""
        row["reference_no"] = record.get("contract_no") or ""
    row["customer_country"] = " / ".join(str(v) for v in (row.get("customer_name"), row.get("country")) if v)
    row.setdefault("lineage", {})["mapping_rule"] = (
        "华康A 2026-09-04统一排期：客户专属字段按表头映射，三表分别在取消单前追加。"
    )
    for field in EXTRA_HEADERS[customer_code]:
        if row.get(field) not in (None, ""):
            row["lineage"][field] = "固定运输规则 40'YT" if customer_code == "green-toys" and field == "transportation_mode" else f"{row.get('source_po_file_name', 'PO')} · {field}"
    return row


def round_cartons(row):
    from app.services.customer_order_unified import _number, _number_text
    quantity, units = _number(row.get("quantity")), _number(row.get("units_per_carton"))
    if quantity is not None and units is not None and units > 0:
        row["carton_count"] = _number_text((quantity / units).to_integral_value(rounding=ROUND_CEILING))


def prepare_360_version(row, history):
    """Business notes belong to this Release/item, not a PO revision."""
    from app.services.customer_order_unified import _text, _issue, _refresh_status
    key = identity('360', row.get('reference_no'), row.get('product_no'))
    if not _text(row.get('product_version')) and all(key):
        candidates = {item.extras['product_version'] for item in history
                      if identity('360', item.reference_no, item.product_no) == key
                      and item.extras.get('product_version')}
        if len(candidates) == 1:
            row['product_version'] = candidates.pop()
            row.setdefault('lineage', {})['product_version'] = 'ITEM表 · 同Release及货号的唯一版本备注'
    row.setdefault('product_version', '')
    row.setdefault('lineage', {}).setdefault('product_version', '业务补填版本备注；非PO REVISION，不按货号推断')
    issue = _issue('warning', 'huakang_a_360_version_note', 'product_version',
                   '版本为业务备注（如两个装、语言版本、拆单说明），可选补填；不填写则保留同单原值或空白。', can_skip=True)
    issue.update(skip_key=f"{row['id']}|huakang_a_360_version_note|product_version",
                 can_edit=True, edit_field='product_version', edit_label='版本', edit_input_type='text')
    row['issues'].append(issue)
    _refresh_status(row)


def ensure_360_version_columns(workbook):
    """Append missing columns without moving historical cells or formulas."""
    from app.services.customer_order_unified import _normalized_header, CustomerOrderUnifiedError
    result = {}
    for name in ('ITEM表', '接单表', '正单评审表'):
        sheet = workbook[name]
        common_end = 28 if name == 'ITEM表' else 19
        found = [c.column for c in sheet[3] if c.column > common_end
                 and _normalized_header(c.value) == '版本']
        if len(found) > 1:
            raise CustomerOrderUnifiedError(f'{name} 的“版本”列重复，请保留唯一表头后重试，不能猜测写入位置')
        column = found[0] if found else sheet.max_column + 1
        if not found:
            header = sheet.cell(3, column)
            header._style = copy(sheet.cell(3, common_end)._style)
            header.value = '版本'
            alignment = copy(header.alignment)
            alignment.wrap_text = True
            header.alignment = alignment
            sheet.column_dimensions[get_column_letter(column)].width = 36
        result[name] = column
    return result


def write_360_version(sheet, row_number, column, row):
    from app.services.customer_order_unified import _text, _fit_generated_text
    cell = sheet.cell(row_number, column)
    if not cell.has_style:
        cell._style = copy(sheet.cell(row_number, 10 if sheet.title == 'ITEM表' else 9)._style)
    alignment = copy(cell.alignment)
    alignment.wrap_text = True
    cell.alignment = alignment
    cell.number_format = '@'
    cell.value = _text(row.get('product_version')) or None
    if cell.value is not None:
        cell.data_type = 's'  # Notes starting '=' must never execute as formulas.
    _fit_generated_text(sheet, row_number, {column: cell.value})


def copy_array_formula(value, source_coordinate, target_coordinate):
    moved = copy(value)
    moved.text = Translator(value.text, origin=source_coordinate).translate_formula(target_coordinate)
    moved.ref = Translator("=" + value.ref, origin=source_coordinate).translate_formula(target_coordinate)[1:]
    return moved


def _insert_rows(workbook, target, start, amount):
    from app.services.huaxing_order_legacy.new_order_excel import (
        _rewrite_formula_for_insert, _shift_target_sheet_structures,
    )
    formulas = [
        (sheet, cell.row, cell.column, copy(cell.value))
        for sheet in workbook for cell in sheet._cells.values()
        if cell.data_type == "f"
    ]
    merges = [CellRange(str(r)) for r in target.merged_cells.ranges if r.max_row >= start]
    for merged in merges:
        target.unmerge_cells(str(merged))
    insert_sparse_rows(target, start, amount)
    _shift_target_sheet_structures(target, start, amount, merges)
    for sheet, old_row, column, value in formulas:
        new_row = old_row + amount if sheet is target and old_row >= start else old_row
        formula = value.text if isinstance(value, ArrayFormula) else value
        rewritten = _rewrite_formula_for_insert(
            formula, formula_sheet=sheet.title, target_sheet=target.title,
            formula_row=old_row, insert_row=start, amount=amount,
        )
        if isinstance(value, ArrayFormula):
            value.text = rewritten
            if sheet is target:
                ref = CellRange(value.ref)
                if ref.min_row >= start:
                    ref.shift(row_shift=amount)
                elif ref.max_row >= start:
                    ref.max_row += amount
                value.ref = str(ref)
        else:
            value = rewritten
        sheet.cell(new_row, column).value = value


def output_slots(workbook, count, *, sheet_counts=None, top_order_area=False):
    from app.services.customer_order_unified import (
        SHEETS, _marker_row, _copy_template_row, CustomerOrderUnifiedError,
    )
    result = {}
    for name, count in (sheet_counts if sheet_counts is not None else dict.fromkeys(SHEETS, count)).items():
        sheet = workbook[name]
        is_total = False
        if top_order_area and count:
            from app.services.customer_order_top_area import boundary
            marker, is_total = boundary(sheet)
            if marker <= 4:
                raise CustomerOrderUnifiedError(f"{name} 顶部正单区缺少明细模板行，请保留至少一行明细格式后重试")
        else:
            marker = _marker_row(sheet)
        # Any literal value is user-owned, including notes without order keys.
        occupied = [r for r in range(4, marker) if any(
            c.value not in (None, "") and c.data_type != "f" for c in sheet[r]
        )]
        start = max(occupied, default=3) + 1
        template_row = next((r for r in range(start - 1, 3, -1) if any(
            c.data_type == "f" for c in sheet[r]
        )), max(4, start - 1) if top_order_area else 4)
        if name != "ITEM表":
            template_row = next((r for r in range(start, marker) if sum(
                c.data_type == "f" for c in sheet[r]
            ) >= 8), template_row)
        free = marker - start
        inserted = max(0, count - free)
        if inserted:
            if top_order_area:
                from app.services.customer_order_top_area import insert_rows
                insert_rows(workbook, sheet, marker, inserted)
            else:
                _insert_rows(workbook, sheet, marker, inserted)
        slots = list(range(start, start + count))
        # Existing reserved formulas are authoritative. Only new rows need a copy.
        for row in slots:
            if row >= marker:
                _copy_template_row(sheet, template_row, row)
            # Some 360 reserved rows have styles but no lookup formulas.
            # Fill only missing formulas, retaining every existing template formula.
            for source in list(sheet[template_row]):
                target = sheet.cell(row, source.column)
                if target.value is None and source.data_type == "f":
                    target._style = copy(source._style)
                    target.value = (
                        copy_array_formula(source.value, source.coordinate, target.coordinate)
                        if isinstance(source.value, ArrayFormula)
                        else Translator(source.value, origin=source.coordinate).translate_formula(target.coordinate)
                    )
        if is_total and slots:
            from app.services.customer_order_top_area import extend_total
            extend_total(sheet, marker + inserted, slots[-1])
        result[name] = slots
    return result


def write_extras(sheet, row_number, row, customer_code):
    from app.services.customer_order_unified import _number, _excel_value, _fit_generated_text
    columns = extra_columns(sheet, customer_code)
    for field, column in columns.items():
        if field in {"amount_usd", "unit_price_hkd", "amount_hkd"}:
            continue
        value = row.get(field)
        if field == "unit_price_usd":
            value = _excel_value(_number(value))
        sheet.cell(row_number, column).value = value if value not in (None, "") else None
        if field == "shipping_item":
            cell = sheet.cell(row_number, column)
            if not cell.has_style:
                cell._style = copy(sheet.cell(row_number, 8)._style)
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            cell.alignment = alignment
            cell.number_format = '@'
            if value not in (None, ""):
                cell.value = str(value)
                cell.data_type = 's'
                _fit_generated_text(sheet, row_number, {column: cell.value})
    if customer_code in {"green-toys", "headstart"}:
        r = row_number
        usd = _number(row.get("unit_price_usd"))
        hkd = _number(row.get("unit_price_hkd"))
        usd_cell = f"{get_column_letter(columns['unit_price_usd'])}{r}" if 'unit_price_usd' in columns else None
        hkd_cell = f"{get_column_letter(columns['unit_price_hkd'])}{r}" if 'unit_price_hkd' in columns else None
        if 'unit_price_hkd' in columns:
            sheet.cell(r, columns['unit_price_hkd']).value = (
                f'=IF({usd_cell}="","",ROUND({usd_cell}*7.8,4))'
                if usd_cell and usd is not None and hkd == (usd * Decimal("7.8")).quantize(Decimal("0.0001"))
                else _excel_value(hkd)
            )
        for field, price_cell, price in (("amount_usd", usd_cell, usd), ("amount_hkd", hkd_cell, hkd)):
            if field in columns:
                quantity = _number(row.get("quantity"))
                sheet.cell(r, columns[field]).value = (
                    f'=IF(OR(K{r}="",{price_cell}=""),"",K{r}*{price_cell})'
                    if price_cell else _excel_value(quantity * price) if quantity is not None and price is not None else None
                )
