"""Refile 360 schedule rows by their existing inspection date on an export copy.

This is a workbook layout operation, not evidence of inspection or shipment.
The sparse row permutation models cut/paste and retains each order exactly once.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from copy import copy, deepcopy
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl.formula import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils.datetime import from_excel, to_excel
from openpyxl.utils.cell import column_index_from_string
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.formatting.formatting import ConditionalFormattingList


SHEETS = ('ITEM表', '接单表', '正单评审表')
COMPLETED = '已完成（按验货日期）'
NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
CELL_REF = re.compile(r"(?:(?P<sheet>'(?:[^']|'')+'|[^ !]+)!)?(?P<c1>\$?[A-Z]{1,3})(?P<r1>\$?\d+)(?::(?P<c2>\$?[A-Z]{1,3})(?P<r2>\$?\d+))?$")


def _error(message):
    from app.services.customer_order_unified import CustomerOrderUnifiedError
    raise CustomerOrderUnifiedError('360 验货日期整理：' + message)


def _text(value):
    return str(value or '').strip()


def _identity(sheet, row):
    ref, product = (5, 8) if sheet.title == 'ITEM表' else (4, 7)
    return tuple(_text(getattr(sheet._cells.get((row, col)), 'value', None)).casefold() for col in (ref, product))


def _plain_date(value, epoch):
    from app.services.customer_order_360_dates import parse_inspection_date
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            value = from_excel(value, epoch=epoch)
        except (ValueError, OverflowError):
            return None
    return parse_inspection_date(value)


def _cache(content):
    """Read only cached formula dates; never trust a cache after formula editing."""
    result = {'dates': {}, 'cells': {}}
    if not content or not content.startswith(b'PK'):
        return result
    with ZipFile(BytesIO(content)) as archive:
        strings = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            strings = [''.join(node.itertext()) for node in ET.fromstring(archive.read('xl/sharedStrings.xml')).findall('m:si', NS)]
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        rels = {r.attrib['Id']: r.attrib['Target'] for r in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        for node in workbook.findall('m:sheets/m:sheet', NS):
            name = node.attrib['name']
            if name not in SHEETS:
                continue
            target = rels[node.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
            target = target.lstrip('/') if target.startswith('/') else 'xl/' + target
            root = ET.fromstring(archive.read(target))
            for cell in root.findall('m:sheetData/m:row/m:c', NS):
                formula, value = cell.find('m:f', NS), cell.find('m:v', NS)
                key = (name, cell.attrib['r'])
                if formula is not None and formula.text:
                    result['cells'][key] = ('f', '=' + formula.text)
                elif cell.attrib.get('t') == 's' and value is not None:
                    result['cells'][key] = ('s', strings[int(value.text)])
                elif cell.attrib.get('t') == 'inlineStr':
                    result['cells'][key] = ('s', ''.join(cell.find('m:is', NS).itertext()))
                elif value is not None and value.text:
                    try:
                        result['cells'][key] = ('n', Decimal(value.text))
                    except InvalidOperation:
                        result['cells'][key] = ('s', value.text)
                if formula is not None and formula.text and value is not None and value.text:
                    try:
                        result['dates'][key] = ('=' + formula.text, float(value.text))
                    except ValueError:
                        pass
    return result


def _signature(cell, epoch):
    if cell is None or cell.value in (None, ''):
        return None
    value = cell.value.text if isinstance(cell.value, ArrayFormula) else cell.value
    if cell.data_type == 'f':
        return ('f', value)
    if isinstance(value, (datetime, date)):
        value = to_excel(value, epoch=epoch)
    if isinstance(value, bool):
        value = int(value)
    if isinstance(value, (int, float, Decimal)):
        return ('n', Decimal(str(value)))
    return ('s', value)


def _dependencies_unchanged(workbook, sheet_name, formula, cached, seen=None):
    """A matching formula alone does not prove its saved result is current."""
    seen = set(seen or ())
    marker = (sheet_name, formula)
    if marker in seen or len(seen) > 30:
        return False
    seen.add(marker)
    functions = {'IF', 'OR', 'AND', 'COUNTIFS', 'COUNTIF', 'INDEX', 'SUMPRODUCT', 'ROW', 'LEN', 'IFERROR', 'IFNA', 'DATE'}
    for token in Tokenizer(formula).items:
        if token.type == 'FUNC' and token.subtype == 'OPEN' and token.value[:-1].upper() not in functions:
            return False
        if token.type != 'OPERAND' or token.subtype != 'RANGE':
            continue
        ref = CELL_REF.fullmatch(token.value)
        if not ref:
            return False  # Names/external/whole-column references are not proven.
        name = (ref.group('sheet') or sheet_name).strip("'").replace("''", "'")
        if name not in workbook.sheetnames:
            return False
        target = workbook[name]
        first = int(ref.group('r1').replace('$', ''))
        last = int((ref.group('r2') or ref.group('r1')).replace('$', ''))
        left = column_index_from_string(ref.group('c1').replace('$', ''))
        right = column_index_from_string((ref.group('c2') or ref.group('c1')).replace('$', ''))
        coordinates = {cell.coordinate for (r, c), cell in target._cells.items() if first <= r <= last and left <= c <= right and cell.value not in (None, '')}
        for source_name, coordinate in cached['cells']:
            if source_name != name:
                continue
            match = CELL_REF.fullmatch(coordinate)
            row = int(match.group('r1'))
            col = column_index_from_string(match.group('c1'))
            if first <= row <= last and left <= col <= right:
                coordinates.add(coordinate)
        for coordinate in coordinates:
            match = CELL_REF.fullmatch(coordinate)
            cell = target._cells.get((int(match.group('r1')), column_index_from_string(match.group('c1'))))
            current = _signature(cell, workbook.epoch)
            original = cached['cells'].get((name, coordinate))
            if current != original and not (current is None and original == ('s', '')):
                return False
            if current and current[0] == 'f' and not _dependencies_unchanged(workbook, name, current[1], cached, seen):
                return False
    return True


def _date_value(workbook, sheet, coordinate, cached, visited=None):
    visited = set(visited or ())
    key = (sheet.title, coordinate)
    if key in visited or len(visited) > 20:
        return None
    visited.add(key)
    cell = sheet[coordinate]
    value = cell.value.text if isinstance(cell.value, ArrayFormula) else cell.value
    if cell.data_type != 'f':
        return _plain_date(value, workbook.epoch)
    # Re-evaluate explicit arithmetic/direct references after openpyxl removed
    # formula caches. Never derive an inspection date from a separate FCD field.
    match = re.fullmatch(r"=(?:(?:'([^']+)'|([^!]+))!)?(\$?[A-Z]+\$?\d+)(?:\s*([+-])\s*(\d+))?", value or '')
    if match:
        name = match.group(1) or match.group(2) or sheet.title
        if name not in workbook.sheetnames:
            return None
        base = _date_value(workbook, workbook[name], match.group(3).replace('$', ''), cached, visited)
        try:
            return base + timedelta(days=int(match.group(5)) * (1 if match.group(4) == '+' else -1)) if base and match.group(5) else base
        except OverflowError:
            return None
    # The unified exporter uses this exact blank-preserving direct link.
    match = re.fullmatch(r'=IF\(LEN\((.+)\)=0,"",\1\)', value or '')
    if match:
        ref = CELL_REF.fullmatch(match.group(1))
        if ref and not ref.group('r2'):
            name = (ref.group('sheet') or sheet.title).strip("'").replace("''", "'")
            if name in workbook.sheetnames:
                return _date_value(workbook, workbook[name], ref.group('c1').replace('$', '') + ref.group('r1').replace('$', ''), cached, visited)
    if sheet.title != 'ITEM表':
        recognized, result = _lookup_date(workbook, sheet, cell.row, value, cached, visited)
        if recognized:
            return result
    if key in cached['dates'] and cached['dates'][key][0] == value and _dependencies_unchanged(workbook, sheet.title, value, cached):
        return _plain_date(cached['dates'][key][1], workbook.epoch)
    return None


def _lookup_date(workbook, sheet, row, formula, cached, visited):
    """Evaluate only the exact identity lookup used by the original 360 file.

    Matching an identity is not sufficient to replace an arbitrary formula.
    The bounded INDEX/SUMPRODUCT expression must itself prove that identity
    and the Y date column, with exactly one source row in its lookup range.
    """
    value = re.sub(r'\s+', '', formula or '').replace('$', '').replace("'ITEM表'!", 'ITEM表!')
    bounds = re.search(r'ITEM表!Y(\d+):Y(\d+)', value)
    if not bounds:
        return False, None
    first, last = map(int, bounds.groups())
    refs, products, dates = (f'ITEM表!{col}{first}:{col}{last}' for col in ('E', 'H', 'Y'))
    lookup = f'INDEX({dates},SUMPRODUCT(({refs}=D{row})*({products}=G{row})*(ROW({refs})-ROW(ITEM表!E{first})+1)))'
    guarded = (f'IF(OR(LEN(D{row})=0,LEN(G{row})=0,COUNTIFS({refs},D{row},{products},G{row})<>1),"",'
               f'IFERROR(IF(LEN({lookup})=0,"",{lookup}),""))')
    if value not in {'=' + lookup, '=' + guarded}:
        return False, None
    item = workbook['ITEM表']
    key = _identity(sheet, row)
    matches = [r for r in _rows(item, first, last + 1) if all(key) and _identity(item, r) == key]
    return True, (_date_value(workbook, item, f'Y{matches[0]}', cached, visited) if len(matches) == 1 else None)


def _sections(sheet):
    markers = [(c.row, _text(c.value)) for c in sheet._cells.values()
               if c.column == 1 and c.data_type != 'f' and isinstance(c.value, str) and
               not all(_identity(sheet, c.row)) and
               (_text(c.value) in {'取消单', '已走货订单', '已走货', '已完成'} or _text(c.value) == COMPLETED)]
    completed = sorted({r for r, label in markers if label.startswith('已完成')})
    if len(completed) > 1:
        _error(f'{sheet.title} 存在多个已完成分区，请先核对分区标题。')
    stop = min((r for r, _ in markers), default=sheet.max_row + 1)
    start = completed[0] if completed else None
    end = min((r for r, _ in markers if start is not None and r > start), default=1_048_577)
    return stop, start, end


def _rows(sheet, first, last):
    return sorted({c.row for c in sheet._cells.values() if first <= c.row < last and c.value not in (None, '')})


def _simulated(sheet, row):
    return any(c.row == row and isinstance(c.value, str) and '模拟数据' in c.value for c in sheet._cells.values())


def _segments(first, last, mapping):
    """(original start, original end, destination start) for a row permutation."""
    start = first
    for row in sorted(r for r in mapping if first <= r <= last):
        if start < row:
            yield start, row - 1, start
        yield row, row, mapping[row]
        start = row + 1
    if start <= last:
        yield start, last, start


def _range_bounds(first, last, mapping):
    spans = sorted((dest, dest + end - start) for start, end, dest in _segments(first, last, mapping))
    merged = []
    for lo, hi in spans:
        if merged and lo <= merged[-1][1] + 1:
            merged[-1] = (merged[-1][0], max(hi, merged[-1][1]))
        else:
            merged.append((lo, hi))
    if len(merged) != 1:
        _error('移动涉及不能保持连续引用的公式范围，请先核对该公式；本次没有生成文件。')
    return merged[0]


def _reference_parts(match, formula_sheet):
    name = (match.group('sheet') or formula_sheet).strip("'").replace("''", "'")
    first = int(match.group('r1').replace('$', ''))
    last = int((match.group('r2') or match.group('r1')).replace('$', ''))
    return name, first, last


def _row_range_baselines(tokens, formula_sheet, maps):
    """Preserve ROW(range)-ROW(range_start)+1 as a positional lookup index.

    The scalar is a range origin, not an order-cell link. Moving it separately
    would silently change every INDEX result when the first order is moved.
    Other affected ROW expressions have no safe inferred meaning and fail.
    """
    significant = [(i, token) for i, token in enumerate(tokens) if token.type != 'WHITE-SPACE']
    overrides, recognized = {}, set()
    for pos in range(len(significant) - 8):
        window = significant[pos:pos + 9]
        values = [t.value.upper() for _, t in window]
        if not (values[0] == values[4] == 'ROW(' and values[2] == values[6] == ')'
                and values[3] == '-' and values[7:] == ['+', '1']):
            continue
        area, anchor = (CELL_REF.fullmatch(window[i][1].value) for i in (1, 5))
        if not area or not anchor or not area.group('r2') or anchor.group('r2'):
            continue
        name, first, last = _reference_parts(area, formula_sheet)
        anchor_name, anchor_row, _ = _reference_parts(anchor, formula_sheet)
        # ROW ignores columns: real review-price formulas use D-range / C-origin.
        if (name, first) != (anchor_name, anchor_row):
            continue
        new_first, _ = _range_bounds(first, last, maps.get(name, {}))
        prefix = anchor.group('sheet') + '!' if anchor.group('sheet') else ''
        overrides[window[5][0]] = prefix + anchor.group('c1') + ('$' if anchor.group('r1').startswith('$') else '') + str(new_first)
        recognized.update((window[0][0], window[4][0]))
    for pos, (index, token) in enumerate(significant):
        if token.value.upper() != 'ROW(' or index in recognized:
            continue
        operand = significant[pos + 1][1].value if pos + 1 < len(significant) else ''
        ref = CELL_REF.fullmatch(operand)
        if ref:
            name, first, last = _reference_parts(ref, formula_sheet)
            affected = any(first <= r <= last for r in maps.get(name, {}))
        else:
            affected = any(maps.values())
        if affected:
            _error('存在无法安全保持计算含义的 ROW 公式，请先核对该公式；本次没有生成文件。')
    return overrides


def _formula(value, formula_sheet, maps):
    if not isinstance(value, str) or not value:
        return value
    prefix = value.startswith('=')
    tokenizer = Tokenizer(value if prefix else '=' + value)
    baselines = _row_range_baselines(tokenizer.items, formula_sheet, maps)
    for index, token in enumerate(tokenizer.items):
        if index in baselines:
            token.value = baselines[index]
            continue
        if token.type != 'OPERAND' or token.subtype != 'RANGE':
            continue
        match = CELL_REF.fullmatch(token.value)
        if not match:
            continue
        sheet_part = match.group('sheet')
        name = sheet_part.strip("'").replace("''", "'") if sheet_part else formula_sheet
        mapping = maps.get(name, {})
        first = int(match.group('r1').replace('$', ''))
        last = int(match.group('r2').replace('$', '')) if match.group('r2') else first
        if not any(first <= r <= last for r in mapping):
            continue
        new_first, new_last = _range_bounds(first, last, mapping)
        start = ('$' if match.group('r1').startswith('$') else '') + str(new_first)
        result = (sheet_part + '!' if sheet_part else '') + match.group('c1') + start
        if match.group('r2'):
            result += ':' + match.group('c2') + ('$' if match.group('r2').startswith('$') else '') + str(new_last)
        token.value = result
    rendered = tokenizer.render()
    return rendered if prefix else rendered[1:]


def _mapped_range(value, mapping):
    area = CellRange(str(value))
    first, last = _range_bounds(area.min_row, area.max_row, mapping)
    area.min_row, area.max_row = first, last
    return str(area)


def _prepare_structures(workbook, maps):
    """Validate and compute everything before touching any cell or row."""
    formulas, structures = [], {}
    for sheet in workbook:
        mapping = maps.get(sheet.title, {})
        for cell in sheet._cells.values():
            if cell.data_type != 'f':
                continue
            value = copy(cell.value)
            if isinstance(value, ArrayFormula):
                value.text = _formula(value.text, sheet.title, maps)
                value.ref = _mapped_range(value.ref, mapping)
            else:
                value = _formula(value, sheet.title, maps)
            formulas.append((sheet.title, mapping.get(cell.row, cell.row), cell.column, value))
        merges = [_mapped_range(r, mapping) for r in sheet.merged_cells.ranges]
        for table in sheet.tables.values():
            area = CellRange(table.ref)
            if any(area.min_row <= r <= area.max_row for r in mapping):
                _error(f'{sheet.title} 的Excel表格覆盖待移动订单，请先将该表格区域转换为普通区域。')
        if mapping and (sheet._images or sheet._charts):
            _error(f'{sheet.title} 含图形对象，当前不能安全迁移整行，请先核对图形锚点。')
        conditional = ConditionalFormattingList()
        for cf, rules in sheet.conditional_formatting._cf_rules.items():
            base = min(cf.sqref.ranges, key=lambda r: (r.min_row, r.min_col))
            origin = f'{base.min_col_letter}{base.min_row}' if hasattr(base, 'min_col_letter') else sheet.cell(base.min_row, base.min_col).coordinate
            for area in cf.sqref.ranges:
                for old_first, old_last, new_first in _segments(area.min_row, area.max_row, mapping):
                    moved = CellRange(min_col=area.min_col, max_col=area.max_col, min_row=new_first, max_row=new_first + old_last - old_first)
                    old_anchor = sheet.cell(old_first, area.min_col).coordinate
                    for rule in rules:
                        cloned = deepcopy(rule)
                        cloned.formula = [_formula(Translator(f if f.startswith('=') else '=' + f, origin=origin).translate_formula(old_anchor), sheet.title, maps).lstrip('=') for f in (rule.formula or [])]
                        conditional.add(str(moved), cloned)
        validations = []
        for validation in sheet.data_validations.dataValidation:
            if not validation.sqref.ranges:
                continue
            base = min(validation.sqref.ranges, key=lambda r: (r.min_row, r.min_col))
            origin = sheet.cell(base.min_row, base.min_col).coordinate
            for area in validation.sqref.ranges:
                for old_first, old_last, new_first in _segments(area.min_row, area.max_row, mapping):
                    cloned = copy(validation)
                    cloned.sqref = str(CellRange(min_col=area.min_col, max_col=area.max_col, min_row=new_first, max_row=new_first + old_last - old_first))
                    for field in ('formula1', 'formula2'):
                        value = getattr(validation, field)
                        if isinstance(value, str) and value:
                            value = Translator(value if value.startswith('=') else '=' + value, origin=origin).translate_formula(sheet.cell(old_first, area.min_col).coordinate)
                            setattr(cloned, field, _formula(value, sheet.title, maps).lstrip('='))
                    validations.append(cloned)
        structures[sheet.title] = (merges, conditional, validations)
    names = [(name, _formula(name.attr_text, '', maps)) for name in workbook.defined_names.values() if name.attr_text]
    for sheet in workbook:
        names.extend((name, _formula(name.attr_text, sheet.title, maps)) for name in sheet.defined_names.values() if name.attr_text)
    return formulas, structures, names


def organize(workbook, *, source_content=None, today=None):
    from app.services.customer_order_360_dates import shanghai_today
    today = today or shanghai_today()
    cached = _cache(source_content)
    result = {'as_of_date': today.isoformat(), 'sheets': {}, 'warnings': []}
    item = workbook['ITEM表']
    dates_by_key = defaultdict(set)
    simulated_keys = set()
    for row in _rows(item, 4, item.max_row + 1):
        key = _identity(item, row)
        if all(key) and _simulated(item, row):
            simulated_keys.add(key)
    sections, participants, resolved = {}, {}, {}
    grouped = defaultdict(lambda: defaultdict(list))
    for name in SHEETS:
        sheet = workbook[name]
        stop, completed, end = sections[name] = _sections(sheet)
        active = _rows(sheet, 4, stop)
        archived = _rows(sheet, completed + 1, end) if completed else []
        participants[name] = (active, archived)
        for row in active + archived:
            key = _identity(sheet, row)
            if not all(key) or key in simulated_keys or _simulated(sheet, row):
                continue
            coordinate = ('Y' if name == 'ITEM表' else 'R') + str(row)
            value = resolved[name, row] = _date_value(workbook, sheet, coordinate, cached)
            grouped[key][name].append(value)
            if name == 'ITEM表':
                dates_by_key[key].add(value)
    conflicts = set()
    for key, by_sheet in grouped.items():
        groups = list(by_sheet.values())
        for pos, left in enumerate(groups):
            for right in groups[pos + 1:]:
                # Equal-size groups can be compared without inventing a batch
                # pairing. Blanks are unknown; the union of known-date counts
                # must fit in the group's actual row count. Two legitimate
                # batches with different dates therefore remain independent.
                if len(left) == len(right):
                    known = Counter(d for d in left if d) | Counter(d for d in right if d)
                    if sum(known.values()) > len(left):
                        conflicts.add(key)
        if key in conflicts:
            result['warnings'].append(f'订单 {key[0]} / {key[1]} 在三表对应记录的验货日期存在冲突，放在待处理区域核对。')
    maps, headers = {}, {}
    for name in SHEETS:
        sheet = workbook[name]
        stop, completed, end = sections[name]
        active, archived = participants[name]
        outgoing, returning = [], []
        for row in active + archived:
            key = _identity(sheet, row)
            if not all(key) or key in simulated_keys or _simulated(sheet, row):
                continue
            coordinate = ('Y' if name == 'ITEM表' else 'R') + str(row)
            value = resolved[name, row]
            if key in conflicts:
                if row in archived:
                    returning.append(row)
                continue
            if name != 'ITEM表' and sheet[coordinate].data_type == 'f' and value is None:
                choices = dates_by_key.get(key, set())
                if len(choices) > 1:
                    result['warnings'].append(f'{name} 第{row}行相同订单存在不同验货日期，放在待处理区域核对。')
                    if row in archived:
                        returning.append(row)
                    continue
            if value is None and sheet[coordinate].data_type == 'f':
                result['warnings'].append(f'{name} 第{row}行验货日期公式无可靠结果，放在待处理区域核对。')
                if row in active:
                    continue
            if row in active and value is not None and value < today:
                outgoing.append(row)
            elif row in archived and (value is None or value >= today):
                returning.append(row)
        stats = {'moved_to_completed': len(outgoing), 'moved_to_pending': len(returning)}
        result['sheets'][name] = stats
        if not outgoing and not returning:
            maps[name] = {}
            continue
        if completed is None:
            completed = sheet.max_row + 2
            end = 1_048_577
            headers[name] = completed
        literal_rows = {c.row for c in sheet._cells.values() if c.value not in (None, '') and c.data_type != 'f'}
        free_active = [r for r in range(4, stop) if r in outgoing or r not in literal_rows]
        if len(free_active) < len(returning):
            _error(f'{name} 待处理区域没有足够空行，无法保留取消/走货区位置；请先在取消单前增加空行。')
        free_archive = list(returning)
        next_row = max([completed, *archived]) + 1
        while len(free_archive) < len(outgoing):
            if next_row >= end or next_row > 1_048_576:
                _error(f'{name} 已完成分区空间不足，无法保留后续区域。')
            free_archive.append(next_row)
            next_row += 1
        mapping = dict(zip(outgoing, free_archive)) | dict(zip(returning, free_active))
        # Close the permutation with the displaced blank/template rows.
        blanks = sorted(set(mapping.values()) - set(mapping))
        holes = sorted(set(mapping) - set(mapping.values()))
        mapping.update(zip(blanks, holes))
        maps[name] = {old: new for old, new in mapping.items() if old != new}
    if not any(maps.values()):
        return result
    formulas, structures, names = _prepare_structures(workbook, maps)
    for sheet in workbook:
        mapping = maps.get(sheet.title, {})
        cells = {(mapping.get(row, row), col): cell for (row, col), cell in sheet._cells.items()}
        for (row, col), cell in cells.items():
            cell.row = row
            if cell.hyperlink:
                cell.hyperlink.ref = cell.coordinate
        sheet._cells = cells
        dimensions = {mapping.get(row, row): copy(value) for row, value in sheet.row_dimensions.items()}
        sheet.row_dimensions.clear()
        for row, value in dimensions.items():
            value.index = row
            sheet.row_dimensions[row] = value
        merges, conditional, validations = structures[sheet.title]
        sheet.merged_cells = MultiCellRange(merges)
        sheet.conditional_formatting = conditional
        sheet.data_validations.dataValidation = validations
        if sheet.title in headers:
            row = headers[sheet.title]
            cell = sheet.cell(row, 1, COMPLETED)
            cell._style = copy(sheet.cell(3, 1)._style)
            sheet.row_dimensions[row].height = 24
        if mapping and sheet.print_area:
            areas = []
            for area in sheet._print_area.ranges:
                expanded = CellRange(str(area))
                expanded.max_row = max(expanded.max_row, max(mapping.values()))
                areas.append(str(expanded))
            sheet.print_area = areas
        sheet._current_row = sheet.max_row
    for name, row, column, value in formulas:
        workbook[name].cell(row, column).value = value
    for name, value in names:
        name.attr_text = value
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = 'auto'
    return result
