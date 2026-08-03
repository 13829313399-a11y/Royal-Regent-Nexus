from __future__ import annotations

import csv
import io
import math
import os
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable

import openpyxl
import pdfplumber
from openpyxl.utils.datetime import from_excel


DEFAULT_DATA_ROOT = Path(r"D:\aaaAI系统项目\业务部入排期\华康A\ai")
PROJECT_ROOT = Path(__file__).resolve().parent
PROJECT_DATA_ROOT = PROJECT_ROOT / "data"
DATA_ROOT_ENV = "SPIN_DATA_ROOT"
SCHEDULE_SHEET = "360客排期表"
CONTRACT_PATTERN = re.compile(r"^RL-\d+-\d+$", re.IGNORECASE)
ACTIVE_SCHEDULE_NAME = "系统当前排期.xlsx"
IGNORED_SOURCE_FOLDERS = {"系统提取结果", "_schedule_archive", "__pycache__"}


@dataclass(frozen=True)
class ColumnDef:
    index: int
    key: str
    label: str
    kind: str = "text"


COLUMNS = [
    ColumnDef(1, "plastic_ready", "胶件复期", "date_or_text"),
    ColumnDef(2, "material_ready", "来料复期", "date_or_text"),
    ColumnDef(3, "style_no", "款号", "identifier"),
    ColumnDef(4, "order_date", "入单日期", "date"),
    ColumnDef(5, "contract_no", "正单合同号", "identifier"),
    ColumnDef(6, "item_no", "产品货号", "identifier"),
    ColumnDef(7, "shipping_item", "出货货号", "identifier"),
    ColumnDef(8, "product_name", "产品名称"),
    ColumnDef(9, "version", "版本"),
    ColumnDef(10, "quantity", "PO数量", "number"),
    ColumnDef(11, "order_status", "订单状态"),
    ColumnDef(12, "inspection_date", "验货日期", "date_or_text"),
    ColumnDef(13, "fcd", "走货期 FCD", "date_or_text"),
    ColumnDef(14, "inspection_result", "验货结果"),
    ColumnDef(15, "third_party_inspection", "第三方公证行验货"),
    ColumnDef(16, "carton_status", "箱唛复期", "date_or_text"),
    ColumnDef(17, "label_details", "利宝明细"),
    ColumnDef(18, "country", "走货国家"),
    ColumnDef(19, "customer", "客户名称"),
    ColumnDef(20, "customer_po", "第三方客户 PO NO#", "identifier"),
    ColumnDef(21, "merchandiser", "跟单"),
    ColumnDef(22, "customer_release", "Customer Release No.", "identifier"),
    ColumnDef(23, "carton_qty", "外箱装箱数", "number"),
    ColumnDef(24, "total_cartons", "总箱数", "number"),
    ColumnDef(25, "date_code", "日期码", "identifier"),
    ColumnDef(26, "barcode", "条码", "identifier"),
    ColumnDef(27, "unit_usd", "订单单价 USD", "number"),
    ColumnDef(28, "unit_hkd", "单价 HKD", "number"),
    ColumnDef(29, "total_usd", "总金额 USD", "number"),
    ColumnDef(30, "total_hkd", "总金额 HKD", "number"),
    ColumnDef(31, "factory_unit_hkd", "出厂价 HKD", "number"),
    ColumnDef(32, "factory_total_hkd", "出厂价总金额 HKD", "number"),
    ColumnDef(33, "remark", "备注"),
    ColumnDef(34, "system_date", "系统", "date_or_text"),
    ColumnDef(35, "container_type", "MS Container Type (FCL/LCL)"),
    ColumnDef(36, "consolidated_id", "Consolidated ID", "identifier"),
    ColumnDef(37, "port", "Port of Discharge"),
    ColumnDef(38, "shipment_date", "走货日期", "date_or_text"),
    ColumnDef(39, "truck", "车次"),
    ColumnDef(40, "shipping_method", "走货方式"),
    ColumnDef(41, "invoice_price", "发票单价", "number"),
    ColumnDef(42, "invoice_amount", "发票金额", "number"),
    ColumnDef(43, "invoice_date", "打票日期", "date_or_text"),
    ColumnDef(44, "invoice_no", "发票号码", "identifier"),
]


CSV_FIELDS = [
    ("excel_row", "Excel行"),
    ("status", "系统状态"),
    *[(column.key, column.label) for column in COLUMNS],
]


def data_root() -> Path:
    configured = os.environ.get(DATA_ROOT_ENV)
    if configured:
        return Path(configured).expanduser().resolve()
    if DEFAULT_DATA_ROOT.exists():
        return DEFAULT_DATA_ROOT.resolve()
    return PROJECT_DATA_ROOT.resolve()


def discover_sources(root: Path | None = None) -> dict[str, Any]:
    root = (root or data_root()).resolve()
    if not root.exists():
        raise FileNotFoundError(f"数据目录不存在：{root}")
    workbooks = sorted(root.glob("*.xlsx"), key=lambda path: path.stat().st_mtime, reverse=True)
    preferred = [path for path in workbooks if "排期" in path.name and not path.name.startswith("~$")]
    if not preferred:
        raise FileNotFoundError(f"数据目录内未找到当前有效排期 Excel；首次使用或业务换版时上传一次：{root}")
    pdfs = sorted(
        (
            path
            for path in root.rglob("*.pdf")
            if not any(part in IGNORED_SOURCE_FOLDERS for part in path.relative_to(root).parts)
        ),
        key=lambda path: (path.parent.name, path.name),
    )
    active_schedule = root / ACTIVE_SCHEDULE_NAME
    schedule = active_schedule if active_schedule.is_file() else preferred[0]
    return {
        "root": str(root),
        "schedule": str(schedule),
        "schedule_name": schedule.name,
        "schedule_modified": datetime.fromtimestamp(schedule.stat().st_mtime).isoformat(timespec="seconds"),
        "pdf_count": len(pdfs),
        "pdfs": [str(path) for path in pdfs],
        "groups": sorted({path.parent.name for path in pdfs}),
    }


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return ""
        if float(value).is_integer():
            return str(int(value))
        return format(float(value), ".15g")
    return re.sub(r"\s+", " ", str(value)).strip()


def _number(value: Any) -> int | float | None:
    if value in (None, "") or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().replace(",", "").replace("$", "")
        try:
            number = float(text)
        except ValueError:
            return None
    if math.isnan(number) or math.isinf(number):
        return None
    if number.is_integer():
        return int(number)
    return round(number, 6)


def _date_text(value: Any, keep_unknown: bool = False) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and 20_000 <= float(value) <= 80_000:
        try:
            converted = from_excel(float(value))
            return converted.date().isoformat() if isinstance(converted, datetime) else converted.isoformat()
        except (TypeError, ValueError, OverflowError):
            pass
    text = _clean_text(value)
    match = re.search(r"(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})", text)
    if match:
        try:
            return date(int(match.group(1)), int(match.group(2)), int(match.group(3))).isoformat()
        except ValueError:
            pass
    for pattern in ("%d-%b-%Y", "%d-%B-%Y", "%m/%d/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    return text if keep_unknown else ""


def _cell_value(value: Any, kind: str) -> Any:
    if kind == "number":
        return _number(value)
    if kind == "date":
        return _date_text(value)
    if kind == "date_or_text":
        return _date_text(value, keep_unknown=True)
    return _clean_text(value)


def _iso_date(value: str) -> date | None:
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError):
        return None


def _row_status(row: dict[str, Any], today: date | None = None) -> str:
    today = today or date.today()
    if row.get("shipment_date"):
        return "已出货"
    fcd = _iso_date(row.get("fcd", ""))
    inspection = _iso_date(row.get("inspection_date", ""))
    if not fcd:
        return "待排期"
    if fcd < today:
        return "已逾期"
    if fcd <= date.fromordinal(today.toordinal() + 14):
        return "临近交期"
    if inspection and inspection < today and not row.get("inspection_result"):
        return "验货待确认"
    return "正常"


def load_schedule(path: Path | str | None = None, sheet_name: str = SCHEDULE_SHEET) -> dict[str, Any]:
    source = discover_sources() if path is None else None
    workbook_path = Path(path or source["schedule"])
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"排期表缺少工作表“{sheet_name}”，现有：{', '.join(workbook.sheetnames)}")
    worksheet = workbook[sheet_name]
    rows: list[dict[str, Any]] = []
    for excel_row, values in enumerate(
        worksheet.iter_rows(min_row=4, max_col=max(column.index for column in COLUMNS), values_only=True),
        start=4,
    ):
        contract_no = _clean_text(values[4] if len(values) > 4 else None).upper()
        if not CONTRACT_PATTERN.fullmatch(contract_no):
            continue
        row = {
            column.key: _cell_value(values[column.index - 1], column.kind)
            for column in COLUMNS
        }
        row["contract_no"] = contract_no
        row["excel_row"] = excel_row
        row["status"] = _row_status(row)
        rows.append(row)
    workbook.close()
    return {
        "sheet": sheet_name,
        "source": str(workbook_path),
        "rows": rows,
        "stats": build_stats(rows),
    }


def build_stats(rows: list[dict[str, Any]]) -> dict[str, Any]:
    customers = sorted({row["customer"] for row in rows if row.get("customer")})
    products = sorted({row["product_name"] for row in rows if row.get("product_name")})
    merchandisers = sorted({row["merchandiser"] for row in rows if row.get("merchandiser")})
    countries = sorted({row["country"] for row in rows if row.get("country")})
    statuses = Counter(row["status"] for row in rows)
    group_quantities: defaultdict[str, float] = defaultdict(float)
    customer_quantities: defaultdict[str, float] = defaultdict(float)
    month_quantities: defaultdict[str, float] = defaultdict(float)
    for row in rows:
        quantity = float(row.get("quantity") or 0)
        customer_quantities[row.get("customer") or "未填写"] += quantity
        product_group = _product_group(row.get("product_name", ""))
        group_quantities[product_group] += quantity
        fcd = row.get("fcd", "")
        month = fcd[:7] if re.fullmatch(r"\d{4}-\d{2}-\d{2}", fcd or "") else "未排期"
        month_quantities[month] += quantity
    dated = [_iso_date(row.get("fcd", "")) for row in rows]
    dated = [value for value in dated if value]
    return {
        "row_count": len(rows),
        "customer_count": len(customers),
        "product_count": len(products),
        "country_count": len(countries),
        "total_quantity": _sum(rows, "quantity"),
        "total_cartons": _sum(rows, "total_cartons"),
        "total_usd": _sum(rows, "total_usd"),
        "total_hkd": _sum(rows, "total_hkd"),
        "statuses": dict(sorted(statuses.items())),
        "customers": customers,
        "products": products,
        "merchandisers": merchandisers,
        "countries": countries,
        "fcd_start": min(dated).isoformat() if dated else "",
        "fcd_end": max(dated).isoformat() if dated else "",
        "ship_months": dict(sorted(month_quantities.items())),
        "top_customers": _top_counts(customer_quantities),
        "product_groups": _top_counts(group_quantities, limit=20),
    }


def _product_group(product_name: str) -> str:
    if not product_name:
        return "未填写"
    match = re.search(r"（([^）]+)）", product_name)
    if match:
        return match.group(1)
    return product_name[:26]


def _top_counts(values: dict[str, float], limit: int = 10) -> list[dict[str, Any]]:
    return [
        {"name": name, "quantity": _round_number(quantity)}
        for name, quantity in sorted(values.items(), key=lambda item: item[1], reverse=True)[:limit]
    ]


def _sum(rows: Iterable[dict[str, Any]], field: str) -> int | float:
    return _round_number(sum(float(row.get(field) or 0) for row in rows))


def _round_number(value: float) -> int | float:
    rounded = round(value, 2)
    return int(rounded) if rounded.is_integer() else rounded


@lru_cache(maxsize=64)
def _extract_pdf(path_text: str, modified_ns: int, size: int) -> tuple[str, int]:
    del modified_ns, size
    path = Path(path_text)
    with pdfplumber.open(path) as pdf:
        page_count = len(pdf.pages)
        text = "\n".join((page.extract_text() or "") for page in pdf.pages[:3])
    return text, page_count


def _extract_workbook(path: Path) -> tuple[str, int]:
    """Flatten WPS PDF-to-Excel output while retaining row order."""
    rows: list[str] = []
    suffix = path.suffix.lower()
    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(path)
        for sheet in book.sheets():
            for row_index in range(sheet.nrows):
                values = [
                    _clean_text(sheet.cell_value(row_index, column_index)).replace("\n", " | ")
                    for column_index in range(sheet.ncols)
                ]
                values = [value for value in values if value]
                if values:
                    rows.append("  ".join(values))
        return "\n".join(rows), len(book.sheets())

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                values = [
                    _clean_text(value).replace("\n", " | ")
                    for value in row
                    if _clean_text(value)
                ]
                if values:
                    rows.append("  ".join(values))
        return "\n".join(rows), len(workbook.worksheets)
    finally:
        workbook.close()


def _parse_po_text(path: Path, text: str, page_count: int, root: Path | None = None) -> dict[str, Any]:
    stat = path.stat()
    contract = _first(r"PURCHASE ORDER RELEASE\s*\([^)]+\)\s*(RL-\d+-\d+)", text).upper()
    release_type = _first(r"PURCHASE ORDER RELEASE\s*\(([^)]+)\)", text)
    revision = _first(r"REVISION\s+(\d+)", text)
    item_match = re.search(
        r"(?m)^(\d+)\s+([A-Za-z0-9.]+)\s+(?:Each|PCS|Set)\s+([\d,]+(?:\.\d+)?)\s*$",
        text,
    )
    converted_match = None
    if not item_match:
        converted_match = re.search(
            r"(?m)^(\d+)\s+([A-Za-z0-9.]+)\b(.*?)\s+(?:Each|PCS|Set)\s+([\d,]+(?:\.\d+)?)\s*$",
            text,
        )
        item_match = converted_match
    line_no = item_match.group(1) if item_match else ""
    item_no = item_match.group(2) if item_match else ""
    quantity = _number(item_match.group(4 if converted_match else 3)) if item_match else None
    description = ""
    if converted_match:
        description = re.sub(r"\s*\|\s*", " ", converted_match.group(3))
        description = re.split(
            r"\b(?:EAN Code|ITF Code|Brand Name|Artwork)\b",
            description,
            maxsplit=1,
            flags=re.I,
        )[0]
        description = _clean_text(description)
    elif item_match:
        tail = text[item_match.end() :]
        description = next((line.strip() for line in tail.splitlines() if line.strip()), "")
    inspection = _first(r"Planned Inspection Date\s*:\s*([^\r\n]+?)(?=\s+Factory Commit Date\s*:|$)", text)
    fcd = _first(r"Factory Commit Date\s*:\s*([^\r\n]+)", text)
    root = (root or path.parent).resolve()
    try:
        relative_path = str(path.resolve().relative_to(root))
    except ValueError:
        relative_path = path.name
    parsed = {
        "file_name": path.name,
        "relative_path": relative_path,
        "group": path.parent.name,
        "page_count": page_count,
        "file_size": stat.st_size,
        "contract_no": contract,
        "release_type": release_type,
        "revision": int(revision) if revision else 0,
        "revision_date": _first(r"Revision Date\s*:\s*([^\r\n]+)", text),
        "customer_po": _first(r"Customer PO Number\s*:\s*([^\r\n]+)", text),
        "customer_release": _first(r"Customer Release No\.\s*:\s*([^\r\n]+)", text),
        "product_specialist": _first(r"Product Specialist\s*:\s*([^\r\n]+)", text),
        "contact": _first(r"Our Contact\s*:\s*([^\r\n]+)", text),
        "container_type": _first(r"MS Container Type \(FCL/LCL\)\s*:\s*([^\r\n]+)", text),
        "order_type": _first(r"RL Order Type\s*:\s*([^\r\n]+)", text),
        "line_no": line_no,
        "item_no": item_no,
        "description": description,
        "quantity": quantity,
        "master_carton_qty": _number(_first(r"Master Carton Qty\s*:\s*([\d,.]+)", text)),
        "planned_inspection_date": _date_text(inspection),
        "factory_commit_date": _date_text(fcd),
        "transportation_mode": _first(r"Transportation Mode\s*:\s*([^\r\n]+)", text),
        "port": _first(r"Port of Discharge[ \t]*:[ \t]*([^\r\n]*)", text),
        "parse_ok": bool(contract and item_no and quantity is not None),
    }
    parsed["parse_warnings"] = [
        label
        for field, label in (
            ("contract_no", "未识别合同号"),
            ("item_no", "未识别货号"),
            ("quantity", "未识别数量"),
            ("factory_commit_date", "未识别 FCD"),
        )
        if parsed.get(field) in (None, "")
    ]
    return parsed


def parse_po_pdf(path: Path, root: Path | None = None) -> dict[str, Any]:
    stat = path.stat()
    text, page_count = _extract_pdf(str(path.resolve()), stat.st_mtime_ns, stat.st_size)
    return _parse_po_text(path, text, page_count, root)


def parse_po_file(path: Path, root: Path | None = None) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_po_pdf(path, root)
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        text, sheet_count = _extract_workbook(path)
        return _parse_po_text(path, text, sheet_count, root)
    raise ValueError(f"华康 A PO 不支持 {suffix or '无扩展名'}")


def _first(pattern: str, text: str, flags: int = re.IGNORECASE) -> str:
    match = re.search(pattern, text, flags)
    return re.sub(r"\s+", " ", match.group(1)).strip() if match else ""


def load_pdfs(paths: Iterable[str], root: Path) -> list[dict[str, Any]]:
    return [parse_po_pdf(Path(path), root) for path in paths]


def validate_pdfs(rows: list[dict[str, Any]], pdfs: list[dict[str, Any]]) -> dict[str, Any]:
    schedule_by_contract: defaultdict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        schedule_by_contract[row["contract_no"].upper()].append(row)

    results: list[dict[str, Any]] = []
    matched_contracts: set[str] = set()
    mismatch_count = 0
    missing_schedule_count = 0
    parse_error_count = 0
    field_passes = 0
    field_checks = 0

    for pdf in pdfs:
        contract = pdf.get("contract_no", "").upper()
        schedule_matches = schedule_by_contract.get(contract, [])
        row = schedule_matches[0] if schedule_matches else None
        if not pdf.get("parse_ok"):
            status = "parse_error"
            parse_error_count += 1
            checks: list[dict[str, Any]] = []
        elif not row:
            status = "not_found"
            missing_schedule_count += 1
            checks = []
        else:
            matched_contracts.add(contract)
            checks = [
                _compare("产品货号", pdf.get("item_no"), row.get("item_no"), "item"),
                _compare("PO 数量", pdf.get("quantity"), row.get("quantity"), "number"),
                _compare("客户 PO", pdf.get("customer_po"), row.get("customer_po"), "identifier"),
                _compare("Customer Release", pdf.get("customer_release"), row.get("customer_release"), "identifier"),
                _compare("外箱装箱数", pdf.get("master_carton_qty"), row.get("carton_qty"), "number"),
                _compare("FCD", pdf.get("factory_commit_date"), row.get("fcd"), "date"),
                _compare("柜型", pdf.get("container_type"), row.get("container_type"), "optional_text"),
                _compare("卸货港", pdf.get("port"), row.get("port"), "optional_text"),
            ]
            field_checks += len(checks)
            field_passes += sum(1 for check in checks if check["ok"])
            status = "pass" if all(check["ok"] for check in checks) else "mismatch"
            if status == "mismatch":
                mismatch_count += 1
        results.append(
            {
                "status": status,
                "file_name": pdf.get("file_name"),
                "group": pdf.get("group"),
                "contract_no": contract,
                "pdf": pdf,
                "schedule": row,
                "schedule_match_count": len(schedule_matches),
                "checks": checks,
                "failed_checks": [check for check in checks if not check["ok"]],
            }
        )

    schedule_without_pdf = [
        {
            "contract_no": row["contract_no"],
            "excel_row": row["excel_row"],
            "product_name": row.get("product_name"),
            "quantity": row.get("quantity"),
            "fcd": row.get("fcd"),
        }
        for row in rows
        if row["contract_no"].upper() not in matched_contracts
    ]
    pass_count = sum(1 for result in results if result["status"] == "pass")
    overall = "pass" if pdfs and pass_count == len(pdfs) else "review"
    return {
        "status": overall,
        "pdf_count": len(pdfs),
        "pass_count": pass_count,
        "mismatch_count": mismatch_count,
        "not_found_count": missing_schedule_count,
        "parse_error_count": parse_error_count,
        "field_pass_rate": round((field_passes / field_checks * 100) if field_checks else 0, 1),
        "matched_schedule_count": len(matched_contracts),
        "schedule_without_pdf_count": len(schedule_without_pdf),
        "schedule_without_pdf": schedule_without_pdf,
        "results": results,
    }


def _compare(label: str, pdf_value: Any, excel_value: Any, kind: str) -> dict[str, Any]:
    pdf_display = "" if pdf_value is None else str(pdf_value)
    excel_display = "" if excel_value is None else str(excel_value)
    if kind == "number":
        left = _number(pdf_value)
        right = _number(excel_value)
        ok = left is not None and right is not None and abs(float(left) - float(right)) <= 0.01
    elif kind == "item":
        left_number = _number(pdf_value)
        right_number = _number(excel_value)
        if left_number is not None and right_number is not None:
            ok = abs(float(left_number) - float(right_number)) <= 0.000001
        else:
            ok = bool(_normal_identifier(pdf_value)) and _normal_identifier(pdf_value) == _normal_identifier(excel_value)
    elif kind == "date":
        left = _date_text(pdf_value)
        right = _date_text(excel_value)
        ok = bool(left and right and left == right)
    elif kind == "optional_text":
        left = _normal_text(pdf_value)
        right = _normal_text(excel_value)
        ok = not left or not right or left == right or left in right or right in left
    else:
        left = _normal_identifier(pdf_value)
        right = _normal_identifier(excel_value)
        ok = bool(left and right and left == right)
    return {
        "label": label,
        "pdf": pdf_display,
        "excel": excel_display,
        "ok": ok,
    }


def _normal_identifier(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", _clean_text(value).upper())


def _normal_text(value: Any) -> str:
    return re.sub(r"[^A-Z0-9\u4e00-\u9fff]", "", _clean_text(value).upper())


def run_quality_checks(rows: list[dict[str, Any]], validation: dict[str, Any]) -> list[dict[str, Any]]:
    return [
        _check_required_fields(rows),
        _check_duplicate_contracts(rows),
        _check_carton_math(rows),
        _check_amount_math(rows),
        _check_date_sequence(rows),
        _check_pdf_coverage(validation),
        _check_pdf_fields(validation),
    ]


def _issue(row: dict[str, Any], message: str, field: str = "") -> dict[str, Any]:
    return {
        "excel_row": row.get("excel_row"),
        "contract_no": row.get("contract_no"),
        "field": field,
        "message": message,
    }


def _quality(name: str, issues: list[dict[str, Any]], description: str) -> dict[str, Any]:
    return {
        "name": name,
        "status": "pass" if not issues else "review",
        "issue_count": len(issues),
        "description": description,
        "issues": issues,
    }


def _check_required_fields(rows: list[dict[str, Any]]) -> dict[str, Any]:
    labels = {
        "item_no": "产品货号",
        "product_name": "产品名称",
        "quantity": "PO 数量",
        "fcd": "FCD",
        "customer_po": "客户 PO",
        "customer_release": "Customer Release",
    }
    issues = []
    for row in rows:
        missing = [label for field, label in labels.items() if row.get(field) in (None, "")]
        if missing:
            issues.append(_issue(row, "缺少：" + "、".join(missing)))
    return _quality("必填字段", issues, "检查合同号对应的关键排期字段是否齐全。")


def _check_duplicate_contracts(rows: list[dict[str, Any]]) -> dict[str, Any]:
    counts = Counter(row["contract_no"] for row in rows)
    issues = [
        _issue(row, f"合同号在主表出现 {counts[row['contract_no']]} 次", "contract_no")
        for row in rows
        if counts[row["contract_no"]] > 1
    ]
    return _quality("合同号唯一性", issues, "正单合同号应唯一，避免 PDF 对错排期行。")


def _check_carton_math(rows: list[dict[str, Any]]) -> dict[str, Any]:
    issues = []
    for row in rows:
        quantity = _number(row.get("quantity"))
        carton_qty = _number(row.get("carton_qty"))
        cartons = _number(row.get("total_cartons"))
        if quantity is None or carton_qty in (None, 0) or cartons is None:
            continue
        expected = math.ceil(float(quantity) / float(carton_qty))
        if abs(expected - float(cartons)) > 0.01:
            issues.append(_issue(row, f"主表 {cartons} 箱，按数量/装箱数应为 {expected} 箱", "total_cartons"))
    return _quality("装箱计算", issues, "按 PO 数量 ÷ 外箱装箱数复算总箱数（不足一箱向上取整）。")


def _check_amount_math(rows: list[dict[str, Any]]) -> dict[str, Any]:
    issues = []
    for row in rows:
        quantity = _number(row.get("quantity"))
        unit_price = _number(row.get("unit_usd"))
        total = _number(row.get("total_usd"))
        if quantity is None or unit_price is None or total is None:
            continue
        expected = round(float(quantity) * float(unit_price), 2)
        if abs(expected - float(total)) > 0.05:
            issues.append(_issue(row, f"主表 USD {total}，数量×单价应为 USD {expected}", "total_usd"))
    return _quality("金额计算", issues, "复算 PO 数量 × 订单单价 USD 与总金额 USD。")


def _check_date_sequence(rows: list[dict[str, Any]]) -> dict[str, Any]:
    issues = []
    for row in rows:
        inspection = _iso_date(row.get("inspection_date", ""))
        fcd = _iso_date(row.get("fcd", ""))
        if inspection and fcd and inspection > fcd:
            issues.append(_issue(row, f"验货日期 {inspection} 晚于 FCD {fcd}", "inspection_date"))
    return _quality("日期先后", issues, "验货日期不应晚于工厂承诺交期 FCD。")


def _check_pdf_coverage(validation: dict[str, Any]) -> dict[str, Any]:
    issues = []
    for result in validation["results"]:
        if result["status"] in {"not_found", "parse_error"}:
            message = "PDF 未在主排期找到合同号" if result["status"] == "not_found" else "PDF 关键字段解析失败"
            issues.append(
                {
                    "excel_row": "",
                    "contract_no": result.get("contract_no"),
                    "field": "pdf",
                    "message": f"{result['file_name']}：{message}",
                }
            )
    return _quality("PDF 覆盖", issues, "每份输入 PO PDF 都应能解析并在排期主表找到合同号。")


def _check_pdf_fields(validation: dict[str, Any]) -> dict[str, Any]:
    issues = []
    for result in validation["results"]:
        schedule = result.get("schedule") or {}
        for check in result.get("failed_checks", []):
            issues.append(
                {
                    "excel_row": schedule.get("excel_row"),
                    "contract_no": result.get("contract_no"),
                    "field": check["label"],
                    "message": f"{check['label']}：PDF“{check['pdf'] or '空'}” / Excel“{check['excel'] or '空'}”",
                }
            )
    return _quality("Excel / PDF 字段", issues, "逐单核对货号、数量、客 PO、Release、装箱数、日期、柜型和卸货港。")


def export_schedule_csv(rows: list[dict[str, Any]]) -> str:
    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow([label for _, label in CSV_FIELDS])
    for row in rows:
        writer.writerow([row.get(field, "") for field, _ in CSV_FIELDS])
    return "\ufeff" + output.getvalue()


def export_audit_csv(validation: dict[str, Any]) -> str:
    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(["PDF文件", "分组", "合同号", "核对状态", "字段", "PDF值", "Excel值", "字段结果"])
    for result in validation["results"]:
        if not result["checks"]:
            writer.writerow(
                [
                    result["file_name"],
                    result["group"],
                    result["contract_no"],
                    result["status"],
                    "",
                    "",
                    "",
                    "",
                ]
            )
        for check in result["checks"]:
            writer.writerow(
                [
                    result["file_name"],
                    result["group"],
                    result["contract_no"],
                    result["status"],
                    check["label"],
                    check["pdf"],
                    check["excel"],
                    "通过" if check["ok"] else "差异",
                ]
            )
    return "\ufeff" + output.getvalue()


def build_payload(root: Path | None = None) -> dict[str, Any]:
    sources = discover_sources(root)
    schedule = load_schedule(Path(sources["schedule"]))
    pdfs = load_pdfs(sources["pdfs"], Path(sources["root"]))
    validation = validate_pdfs(schedule["rows"], pdfs)
    quality = run_quality_checks(schedule["rows"], validation)
    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "sources": sources,
        "schedule": schedule,
        "pdfs": pdfs,
        "validation": validation,
        "quality": quality,
    }
