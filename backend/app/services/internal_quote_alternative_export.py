"""Compare independent packaging alternatives using their frozen export bytes."""
import json
import math
import re

from fastapi import HTTPException
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter, quote_sheetname
from sqlalchemy import select

from app.models.internal_quote import InternalQuoteAlternative, InternalQuoteExportFile
from app.services.internal_quote import _get_quote, _check_revision, _add_audit, ensure_quote_read
from app.services.internal_quote_alternatives import issue_alternative
from app.services.internal_quote_artifacts import _ensure_export_permission, create_controlled_export, digest, safe_file_name
from app.services.internal_quote_calculator import canonical_json
from app.services.internal_quote_document import product_root, lock_documents, ensure_not_archived
from app.services.internal_quote_series_export import combine_series_workbooks
from app.services.internal_quote_workbook_view import WorkbookValues
from app.services.transaction_lock import lock_transaction


def _sales_payload(workbook):
    if "结构化数据" not in workbook.sheetnames:
        raise ValueError("原输出缺少包装结构化资料")
    chunks = [row for row in workbook["结构化数据"].iter_rows(min_row=4, values_only=True)
              if row[0] == "payload" and row[1] == "sales"]
    chunks.sort(key=lambda row: int(row[8]))
    if not chunks or [int(row[8]) for row in chunks] != list(range(1, int(chunks[0][9]) + 1)):
        raise ValueError("原输出的包装资料分片不完整")
    return json.loads("".join(str(row[10]) for row in chunks))


def _prices(workbook):
    # Read the actual frozen workbook's price cells, never today's calculator.
    # Component layouts retain the original total-price sheet as a dependency.
    sheet = workbook["_报价明细计算"] if "_报价明细计算" in workbook.sheetnames else workbook.worksheets[0]
    values = WorkbookValues(workbook)
    result = []
    for row in sheet:
        label = str(row[1].value or "") if len(row) > 1 else ""
        match = re.fullmatch(r"报价[（(]MOQ([\d,]+)(K?)[）)]", label, re.IGNORECASE)
        if not match and label != "报价合计":
            continue
        moq = int(match[1].replace(",", "")) * (1000 if match[2] else 1) if match else values.cell(sheet.title, "R6")
        address = f"D{row[0].row}"
        amount = values.cell(sheet.title, address)
        if isinstance(amount, bool) or not isinstance(amount, (int, float)) or not math.isfinite(amount) or amount < 0:
            raise ValueError("原输出出厂价不可读取")
        if not isinstance(moq, (int, float)) or not math.isfinite(moq) or moq <= 0 or int(moq) != moq:
            raise ValueError("原输出 MOQ 不完整")
        result.append((int(moq), amount, sheet.title, address))
    if not result:
        raise ValueError("原输出未找到可核对的 MOQ 出厂价")
    return result


def _text(cell, value):
    cell.value = str(value or "")
    cell.data_type = "s"  # User names/packaging must never become formulas.


def combine_alternative_workbooks(items):
    def summary(target, source, names, index):
        meta = items[index - 1]
        try:
            sales = _sales_payload(source)
            prices = _prices(source)
        except (ValueError, TypeError, KeyError, IndexError, ZeroDivisionError, OverflowError) as error:
            raise HTTPException(409, f"{meta['label']}：{error}；请保留原文件，复制新版本核对后再输出") from error
        packaging = "；".join(" · ".join(str(row.get(key) or "").strip() for key in ("item", "specification")
                                        if str(row.get(key) or "").strip())
                              for row in sales.get("packaging_materials", []) if isinstance(row, dict)) or "未填写包装说明"
        cartons = "；".join(f"{row.get('item') or '纸箱'}：{row.get('qty_per_carton', '')} PCS/箱"
                            for row in sales.get("cartons", []) if isinstance(row, dict))
        if index == 1:
            sheet = target.create_sheet("方案汇总", 0)
            sheet.append(["多方案报价对比"])
            sheet.append(["报价取自各方案原输出的出厂价（HKD/件）；仅列原文件参与输出的 MOQ，不叠加方案成本。"])
            sheet.append(["详细成本保留各自工作表；本汇总为输出时快照，原历史文件保持不变。"])
            sheet.append(["包装方案", "版本", "包装材料 / 规格", "装箱", "MOQ（PCS）", "出厂价 HKD/件", "详细成本", "原输出时间"])
            widths = [25, 13, 48, 28, 16, 20, 24, 23]
            for col, width in enumerate(widths, 1):
                sheet.column_dimensions[get_column_letter(col)].width = width
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = "C5"
            sheet.print_title_rows = "1:4"
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet["A1"].font = Font(name="微软雅黑", size=18, bold=True, color="0F766E")
            sheet.row_dimensions[1].height = 30
            for row_index in (2, 3):
                sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=8)
                sheet.cell(row_index, 1).alignment = Alignment(wrap_text=True, vertical="center")
                sheet.row_dimensions[row_index].height = 28
            for cell in sheet[4]:
                cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
                cell.fill = PatternFill("solid", fgColor="0F766E")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.row_dimensions[4].height = 32
        sheet = target["方案汇总"]
        for moq, price, price_sheet, address in prices:
            row_index = sheet.max_row + 1
            fields = [meta["scenario"], meta["version"], packaging, cartons, moq, price, "查看详细成本", meta["exported_at"]]
            for column, value in enumerate(fields, 1):
                cell = sheet.cell(row_index, column)
                if column in (5, 6):
                    cell.value = value
                    cell.number_format = "#,##0" if column == 5 else "#,##0.00"
                else:
                    _text(cell, value)
                cell.font = Font(name="微软雅黑", size=11, color="0F766E" if column == 6 else "334155", bold=column == 6)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                cell.fill = PatternFill("solid", fgColor="F0FDFA" if index % 2 else "F8FAFC")
            sheet.cell(row_index, 7).hyperlink = f"#{quote_sheetname(names[source.worksheets[0].title])}!A1"
            # A comment traces the displayed frozen amount to the exact detail cell.
            sheet.cell(row_index, 6).comment = Comment(f"原文件：{meta['file_name']}\n来源：{names[price_sheet]}!{address}\nSHA256：{meta['sha256']}", "Royal Regent Nexus")
            sheet.row_dimensions[row_index].height = max(48, 18 * (1 + len(packaging) // 30))
        sheet.auto_filter.ref = f"A4:H{sheet.max_row}"
        sheet.print_area = f"A1:H{sheet.max_row}"
    return combine_series_workbooks([(item["label"], item["content"]) for item in items], summary_builder=summary)


def export_alternatives(db, quote_id, payload, user, request=None):
    try:
        ids = [row.quote_id for row in payload.versions]
        if len(ids) != len(set(ids)):
            raise HTTPException(400, "请勿重复选择同一方案版本")
        source = _get_quote(db, quote_id)
        quotes = [_get_quote(db, identity) for identity in ids]
        for quote in [source, *quotes]:
            ensure_quote_read(db, user, quote.factory_id)
            _ensure_export_permission(db, quote, user)
        lock_documents(db, [source, *quotes])
        roots = [product_root(db, quote) for quote in [source, *quotes]]
        for identity in sorted({f"{q.factory_id}:{q.batch_id or q.id}" for q in [source, *quotes, *roots]}):
            lock_transaction(db, "internal-quote", identity)
        for identity in sorted({root.id for root in roots}):
            lock_transaction(db, "internal-quote-family", identity)
        db.expire_all()
        source = _get_quote(db, quote_id)
        quotes = [_get_quote(db, identity) for identity in ids]
        root_id = product_root(db, source).id
        for quote in [source, *quotes]:
            ensure_not_archived(db, quote)
            if product_root(db, quote).id != root_id or quote.factory_id != source.factory_id or quote.customer != source.customer:
                raise HTTPException(400, "只能合并同一产品、客户和厂区的包装方案")
        records, items = [], []
        for quote, selected in zip(quotes, payload.versions):
            # Prefer the original immutable final export, including legacy layouts.
            candidates = db.scalars(select(InternalQuoteExportFile).where(
                InternalQuoteExportFile.quote_id == quote.id,
                InternalQuoteExportFile.release_stage.in_(["p4_direct_issued", "p4_final_approved"]),
            ).order_by(InternalQuoteExportFile.exported_at, InternalQuoteExportFile.id)).all()
            record = next((row for row in candidates if
                           json.loads(row.export_manifest_json).get("final_release_revision") == quote.final_release_revision), None)
            if record is not None and (quote.status not in {"fully_approved", "exported"}
                                       or quote.final_release_status not in {"approved", "issued"}):
                raise HTTPException(409, "此方案已有历史输出，但当前版本尚未核准，请先核对版本状态")
            if record is None:
                if quote.status == "exported" or quote.final_release_status == "issued":
                    raise HTTPException(409, "已输出版本缺少匹配的原文件，不能重新生成历史报价")
                _check_revision(quote.header_revision, selected.revision, "报价版本")
                if quote.module_version == "v4":
                    out = issue_alternative(db, quote.id, selected, user, request, commit=False)
                else:
                    if not (quote.status in {"fully_approved", "exported"} and quote.final_release_status == "approved" and quote.final_release_revision > 0):
                        raise HTTPException(409, "历史报价须通过最终审核后才能合并输出")
                    out = create_controlled_export(db, quote.id, user, request, commit=False)
                record = db.get(InternalQuoteExportFile, out.id)
            if digest(record.content) != record.sha256:
                raise HTTPException(409, "原输出文件校验失败，请先核对历史文件")
            entry = db.get(InternalQuoteAlternative, quote.id)
            scenario = entry.scenario_name if entry else "原方案"
            label = f"{scenario} {quote.version_label}"
            items.append(dict(label=label, scenario=scenario, version=quote.version_label, content=record.content,
                              exported_at=record.exported_at, file_name=record.file_name, sha256=record.sha256))
            records.append(record)
        content = combine_alternative_workbooks(items)
        sha256 = digest(content)
        _add_audit(db, source, user, "export_alternatives", detail=canonical_json({
            "sha256": sha256, "versions": [{"quote_id": q.id, "export_id": r.id, "sha256": r.sha256}
                                          for q, r in zip(quotes, records)],
        }), request=request)
        db.commit()
        return content, safe_file_name(f"{source.quote_no}-{source.product_name}-多方案报价.xlsx"), sha256
    except Exception:
        db.rollback()
        raise
