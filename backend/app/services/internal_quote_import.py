from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from typing import Any, Callable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


IMPORT_TYPE_DEPARTMENTS = {
    "mold": "engineering",
    "hardware": "engineering",
    "electronic": "electronic",
    "molding": "molding",
    "painting": "painting",
    "slush": "slush",
    "sewing": "sewing",
    "hair": "hair",
    "assembly": "assembly",
}
MAX_WORKBOOK_ROWS = 2000
MAX_WORKBOOK_COLUMNS = 100
MAX_OOXML_UNCOMPRESSED_SIZE = 120 * 1024 * 1024
MAX_OOXML_ENTRIES = 5000

PAINTING_PROCESSES = (
    ("夹模", "clamp"),
    ("移印", "pad_print"),
    ("UV", "uv"),
    ("散枪", "spray"),
    ("边模", "edge"),
    ("油色", "paint"),
    ("浸油", "dip"),
    ("抹油", "wipe"),
    ("擦PP水", "pp_water"),
)


@dataclass
class EmbeddedWorkbookImage:
    source_row: int
    source_column: int
    file_name: str
    content_type: str
    content: bytes


@dataclass
class ParsedInternalQuoteImport:
    target_department: str
    sheet_name: str
    header_row: int
    row_count: int
    payload_fragment: dict[str, Any]
    warnings: list[str] = field(default_factory=list)
    embedded_images: list[EmbeddedWorkbookImage] = field(default_factory=list)


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def normalized(value: object) -> str:
    return (
        re.sub(r"\s+", "", text(value))
        .replace("（", "(")
        .replace("）", ")")
        .upper()
    )


def number(value: object, default: Decimal | None = None) -> Decimal | None:
    if value is None or value == "":
        return default
    if isinstance(value, Decimal):
        result = value
    elif isinstance(value, (int, float)):
        result = Decimal(str(value))
    else:
        token = re.sub(r"[^\d.\-]", "", text(value).replace(",", ""))
        if not token or token in {"-", ".", "-."}:
            return default
        try:
            result = Decimal(token)
        except InvalidOperation:
            return default
    return result if result.is_finite() else default


def decimal_text(value: Decimal | int | float | str | None) -> str:
    parsed = number(value, Decimal("0")) or Decimal("0")
    return format(parsed.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP), "f")


def precise_decimal_text(value: Decimal | int | float | str | None) -> str:
    """Keep source precision while retaining the form's four-decimal minimum."""
    parsed = number(value, Decimal("0")) or Decimal("0")
    raw = format(parsed, "f")
    integer, dot, fraction = raw.partition(".")
    if not dot:
        return f"{integer}.0000"
    fraction = fraction.rstrip("0")
    return f"{integer}.{fraction.ljust(4, '0')}"


def _validate_ooxml(content: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_OOXML_ENTRIES:
                raise ValueError("Excel 内部文件数量异常，已拒绝解析")
            if sum(item.file_size for item in entries) > MAX_OOXML_UNCOMPRESSED_SIZE:
                raise ValueError("Excel 解压后体积超过限制，已拒绝解析")
            if "[Content_Types].xml" not in archive.namelist():
                raise ValueError("文件不是有效的 xlsx/xlsm 工作簿")
    except zipfile.BadZipFile as error:
        raise ValueError("文件不是有效的 xlsx/xlsm 工作簿") from error


def _xls_workbook_rows(content: bytes) -> list[tuple[str, list[list[object]]]]:
    import xlrd

    try:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True, ragged_rows=True)
        try:
            sheets = []
            for sheet in workbook.sheets():
                rows = []
                for row_index in range(min(sheet.nrows, MAX_WORKBOOK_ROWS)):
                    rows.append([
                        xlrd.error_text_from_code.get(cell.value, "#VALUE!")
                        if cell.ctype == xlrd.XL_CELL_ERROR else cell.value
                        for cell in sheet.row(row_index)[:MAX_WORKBOOK_COLUMNS]
                    ])
                if any(any(text(cell) for cell in row) for row in rows):
                    sheets.append((sheet.name, rows))
        finally:
            workbook.release_resources()
    except Exception as error:
        raise ValueError("无法读取车发部 XLS 报价单，请确认文件有效且未加密") from error
    if not sheets:
        raise ValueError("工作簿为空")
    return sheets


def workbook_rows(content: bytes, *, allow_xls: bool = False) -> list[tuple[str, list[list[object]]]]:
    if allow_xls and content.startswith(b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"):
        return _xls_workbook_rows(content)
    _validate_ooxml(content)
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as error:
        raise ValueError("无法读取 Excel，请确认文件为有效的 xlsx/xlsm 工作簿") from error
    sheets: list[tuple[str, list[list[object]]]] = []
    try:
        for sheet in workbook.worksheets:
            rows = [
                    list(row[:MAX_WORKBOOK_COLUMNS])
                    for row in sheet.iter_rows(
                        min_row=1,
                        max_row=min(sheet.max_row or MAX_WORKBOOK_ROWS, MAX_WORKBOOK_ROWS),
                        values_only=True,
                )
            ]
            if any(any(text(cell) for cell in row) for row in rows):
                sheets.append((sheet.title, rows))
    finally:
        workbook.close()
    if not sheets:
        raise ValueError("工作簿为空")
    return sheets


def workbook_embedded_images(
    content: bytes,
    sheet_name: str,
    *,
    source_column: int,
    min_source_row: int,
) -> list[EmbeddedWorkbookImage]:
    """Extract floating Excel images anchored to a specific worksheet column.

    openpyxl exposes worksheet drawing images separately from cell values.  The
    importer otherwise sees only an empty picture cell; this function preserves
    the actual image bytes and their visual row so the confirmation step can
    attach them to the corresponding mold record.
    """
    try:
        workbook = load_workbook(BytesIO(content), data_only=True, read_only=False)
    except Exception as error:
        raise ValueError("无法读取 Excel 中的嵌入图片") from error
    extracted: list[EmbeddedWorkbookImage] = []
    row_counts: dict[int, int] = {}
    try:
        if sheet_name not in workbook.sheetnames:
            return []
        sheet = workbook[sheet_name]

        def row_height_emu(row_number: int) -> int:
            row_height_points = (
                sheet.row_dimensions[row_number].height
                or sheet.sheet_format.defaultRowHeight
                or 15
            )
            return max(int(float(row_height_points) * 12700), 1)

        def marker_absolute_y(marker: object) -> int:
            marker_row = int(getattr(marker, "row", 0) or 0)
            return (
                sum(row_height_emu(row_number) for row_number in range(1, marker_row + 1))
                + int(getattr(marker, "rowOff", 0) or 0)
            )

        def row_at_y(position_emu: int, last_row: int) -> int:
            cumulative = 0
            for row_number in range(1, max(last_row, 1) + 1):
                cumulative += row_height_emu(row_number)
                if position_emu < cumulative:
                    return row_number
            return max(last_row, 1)

        for image in getattr(sheet, "_images", []):
            anchor = getattr(image, "anchor", None)
            marker = getattr(anchor, "_from", None)
            if marker is None:
                continue
            row = int(marker.row) + 1
            end_marker = getattr(anchor, "to", None)
            if end_marker is not None:
                # Use the picture's vertical center rather than only its top-left
                # anchor. Supplier files may start an image at the bottom of the
                # preceding tall row while the picture visibly occupies the next
                # mold row.
                start_y = marker_absolute_y(marker)
                end_y = marker_absolute_y(end_marker)
                row = row_at_y(
                    (start_y + end_y) // 2,
                    max(sheet.max_row, int(end_marker.row) + 1),
                )
            column = int(marker.col) + 1
            if column != source_column or row < min_source_row:
                continue
            try:
                content_bytes = image._data()
            except Exception:
                continue
            image_format = str(getattr(image, "format", "png") or "png").lower()
            if image_format in {"jpg", "jpeg"}:
                extension, content_type = ".jpg", "image/jpeg"
            elif image_format == "webp":
                extension, content_type = ".webp", "image/webp"
            else:
                extension, content_type = ".png", "image/png"
            row_counts[row] = row_counts.get(row, 0) + 1
            column_letter = get_column_letter(source_column)
            extracted.append(
                EmbeddedWorkbookImage(
                    source_row=row,
                    source_column=column,
                    file_name=f"模具图片-{column_letter}{row}-{row_counts[row]}{extension}",
                    content_type=content_type,
                    content=content_bytes,
                )
            )
    finally:
        workbook.close()
    return extracted


def row_has(row: list[object], *patterns: str) -> bool:
    joined = "|".join(normalized(cell) for cell in row)
    return all(any(normalized(token) in joined for token in pattern.split("|")) for pattern in patterns)


def find_header(
    sheets: list[tuple[str, list[list[object]]]],
    import_type: str,
) -> tuple[str, list[list[object]], int]:
    if import_type == "hair":
        # Never take the stale price from 正式 when this template has 明细.
        detail_sheets = [(name, rows) for name, rows in sheets if re.search(r"明细|明細", name)]
        if not detail_sheets:
            raise ValueError("车发报价必须使用‘明细’工作表中的单价和重量")
        sheets = detail_sheets
    patterns = {
        "mold": (
            "模号|模具编号|MOLDNO|产品名称|ITEMDESCRIPTION|项目内容",
            "名称|模价|总价|材质|材料|AMOUNT|MOULDPRICES",
        ),
        "hardware": ("零件名称", "规格", "用量", "单价"),
        "electronic": ("零件名称", "规格", "用量"),
        "molding": ("模具名称|货名", "啤净重|日产量|预估料重", "材质|用料"),
        "sewing": ("物料名称|布料名称", "裁片部位|部位", "用量|用量/码", "价钱|总价钱"),
        "assembly": ("工序名称|做工名称", "人数"),
        "painting": ("名称|位置", "夹模|移印|UV|散枪|边模|抹油|擦PP水"),
        "slush": ("产品编号|产品编码|货号", "胶件名称|零件名称|产品名称", "材料|材质", "用量|数量"),
        "hair": ("货名|名称", "单价", "重量", "单位"),
    }[import_type]
    for sheet_name, rows in sheets:
        for index, row in enumerate(rows[:40]):
            if row_has(row, *patterns):
                return sheet_name, rows, index
    raise ValueError(
        {
            "mold": "未找到模具编号/名称/价格表头，请使用展兴模具--工模报价表.xlsx格式",
            "hardware": "未找到零件名称/规格/用量/单价表头，请使用五金1.xlsx格式",
            "electronic": "未找到零件名称/规格/用量表头",
            "molding": "未找到注塑的模具名称/啤净重表头或吹气的货名/日产量表头",
            "sewing": "未找到布料名称/部位/用量/价钱表头",
            "assembly": "未找到工序名称或做工名称/人数表头",
            "painting": "未找到名称/位置及喷油工序表头",
            "slush": "未找到产品编号/胶件名称/材料/用量表头",
            "hair": "未在明细页找到货名/港币单价/重量/单位表头，请使用车发部报价单模板",
        }[import_type]
    )


def column_index(header: list[object], aliases: tuple[str, ...]) -> int | None:
    tokens = tuple(normalized(alias) for alias in aliases)
    for index, value in enumerate(header):
        header_token = normalized(value)
        if header_token and any(alias in header_token for alias in tokens):
            return index
    return None


def preferred_column_index(header: list[object], aliases: tuple[str, ...]) -> int | None:
    """Prefer an exact normalized header before falling back to a contains match."""
    tokens = tuple(normalized(alias) for alias in aliases)
    for index, value in enumerate(header):
        if normalized(value) in tokens:
            return index
    return column_index(header, aliases)


def value_at(row: list[object], index: int | None) -> object:
    return row[index] if index is not None and index < len(row) else None


def _next_number(row: list[object], start: int) -> Decimal | None:
    for value in row[start + 1 :]:
        parsed = number(value)
        if parsed is not None:
            return parsed
    return None


def _parse_mold(
    rows: list[list[object]],
    header_index: int,
    *,
    fallback_qty: Decimal,
    sheet_name: str = "",
    embedded_images_by_row: dict[int, list[EmbeddedWorkbookImage]] | None = None,
    **_: object,
) -> tuple[dict[str, Any], int, list[str]]:
    header = rows[header_index]
    def layout_token(index: int) -> str:
        return normalized(value_at(header, index)).replace("'", "").replace("’", "")

    is_water_table_layout = (
        len(header) >= 21
        and "MOLDNO" in layout_token(1)
        and "CHINESENAME" in layout_token(4)
        and "MATL" in layout_token(7)
        and "DIM" in layout_token(11)
        and "TOOLINSERTMATL" in layout_token(13)
        and "GATE" in layout_token(18)
        and "PICTURES" in layout_token(20)
    )
    is_zhanxing_mold_base_layout = (
        len(header) >= 12
        and "ITEMDESCRIPTION" in layout_token(1)
        and "MATERIAL" in layout_token(2)
        and "CAVITIES" in layout_token(3)
        and "MOULDPRICES" in layout_token(7)
    )
    is_zhanxing_split_capacity_machine_layout = (
        is_zhanxing_mold_base_layout
        and len(header) >= 13
        and "产能" in layout_token(9)
        and "机型" in layout_token(10)
        and "重量" in layout_token(11)
        and "REMARK" in layout_token(12)
    )
    is_zhanxing_combined_capacity_machine_layout = (
        is_zhanxing_mold_base_layout
        and (
            "日产能" in layout_token(9)
            or "机台大小" in layout_token(9)
            or ("产能" in layout_token(9) and "机型" in layout_token(9))
        )
    )
    is_zhanxing_mold_layout = (
        is_zhanxing_split_capacity_machine_layout
        or is_zhanxing_combined_capacity_machine_layout
    )
    columns = {
        "mold_no": column_index(header, ("模号", "模具编号", "客人模具编号", "MOLD NO")),
        "name": column_index(header, ("产品名称", "零件名称", "配件名称", "模具名称", "加工内容", "ITEM DESCRIPTION", "项目内容", "DESCRIPTION", "名称")),
        "chinese_name": column_index(header, ("中文名称", "中文名", "CHINESE NAME")),
        "material": column_index(header, ("产品材质", "塑胶原料", "胶料类型", "材质", "材料")),
        "material_type": column_index(header, ("料型", "材料类型", "MAT'L", "MAT’L")),
        "weight": column_index(header, ("零件重量", "克重", "净重", "PART WEIGHT")),
        "cavity": column_index(header, ("出模数", "型腔", "模数", "CAV")),
        "sets": column_index(header, ("套数", "数量/件", "数量", "UP")),
        "price": column_index(header, ("含税模价", "总价", "模价", "TOTAL AMOUNT", "AMOUNT")),
        "machine": column_index(header, ("机型(TON)", "INJECTION MACHINE TYPE", "机型")),
        "target": column_index(header, ("模具预计日啤数", "目标数", "CYCLES/DAY")),
        "mold_base_type": column_index(header, ("模胚类型", "模胚型号", "模胚")),
        "mold_base_material": column_index(header, ("模胚材质", "模仁材质", "TOOL INSERT MAT'L", "TOOL INSERT MAT’L")),
        "structure": column_index(header, ("模具结构", "滑块", "行位", "加工内容")),
        "process": column_index(header, ("工艺", "入水方式", "水口", "GATE")),
        "cycle": column_index(header, ("周期(秒)", "周期", "啤塑周期", "CYCLE TIME", "CYCLE")),
        "mold_size": column_index(header, ("模具尺寸", "模胚尺寸")),
        "mold_specification": column_index(header, ("模具规格", "TOOL INFORMATION DIM", "DIM(HXWXD CM)", "DIM")),
        "color": column_index(header, ("颜色", "COLOR")),
        "image": column_index(header, ("图片", "图  片", "IMAGE")),
        "note": column_index(header, ("备注", "说明", "REMARK")),
    }
    if is_water_table_layout:
        # Water Table 模价表使用固定双行表头。业务确认的字段位置为：
        # B 模号、E 模具名称、H 料型、L 模具尺寸、N 模胚材质、S 工艺、U 图片。
        # 仅覆盖这些明确列，其他字段继续沿用通用表头匹配。
        columns.update({
            "mold_no": 1,
            "name": 4,
            "chinese_name": None,
            "material_type": 7,
            "mold_size": 11,
            "mold_specification": None,
            "mold_base_material": 13,
            "process": 18,
            "image": 20,
        })
    if is_zhanxing_mold_layout:
        # 展兴工模报价表使用固定列。新版为 A–M：
        # A 模号、B 中文名称、C 料型、D 出模数、E 套数、F 图片、
        # G 模具尺寸、H 模价 RMB、I 模胚材质、J 产能、K 机型、
        # L 净重、M 备注。旧版 A–L 的 J 列仍兼容“产能/机型”。
        columns.update({
            "mold_no": 0,
            "name": 1,
            "chinese_name": 1,
            "material": None,
            "material_type": 2,
            "cavity": 3,
            "sets": 4,
            "image": 5,
            "mold_size": 6,
            "mold_specification": None,
            "price": 7,
            "mold_base_material": 8,
            "machine": 10 if is_zhanxing_split_capacity_machine_layout else None,
            "target": 9 if is_zhanxing_split_capacity_machine_layout else None,
            "weight": 11 if is_zhanxing_split_capacity_machine_layout else 10,
            "note": 12 if is_zhanxing_split_capacity_machine_layout else 11,
        })
    warnings: list[str] = []
    molds: list[dict[str, Any]] = []
    last_mold_no = ""
    amortization_candidates: list[Decimal] = []

    explicit_amortization: Decimal | None = None
    for row in rows:
        for index, cell in enumerate(row):
            if re.search(r"分摊(数量|套数)|摊分(数量|套数)", text(cell)):
                explicit_amortization = _next_number(row, index)
                if explicit_amortization and explicit_amortization > 0:
                    break

    for source_row, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        joined = "".join(text(cell) for cell in row)
        if re.search(r"^[一二三四五六七八九十]、.+部分", joined):
            break
        if (
            is_zhanxing_mold_layout
            and re.search(r"工模.*注塑模", normalized(value_at(row, 0)))
        ):
            break
        if re.search(r"^(合计|小计|总计|说明|备注|客户确认|签名)", joined):
            continue
        raw_mold_no = text(value_at(row, columns["mold_no"]))
        name = text(value_at(row, columns["name"]))
        chinese_name = text(value_at(row, columns["chinese_name"]))
        price = number(value_at(row, columns["price"]))
        combined_capacity_machine = (
            text(value_at(row, 9))
            if is_zhanxing_combined_capacity_machine_layout
            else ""
        )
        machine_code = text(value_at(row, columns["machine"]))
        target_output = number(value_at(row, columns["target"]), Decimal("0"))
        if combined_capacity_machine:
            capacity_machine_parts = re.split(
                r"[/／]",
                combined_capacity_machine,
                maxsplit=1,
            )
            target_output = number(capacity_machine_parts[0], Decimal("0"))
            machine_code = (
                text(capacity_machine_parts[1])
                if len(capacity_machine_parts) > 1
                else ""
            )
        continuation_values = (
            name,
            chinese_name,
            text(value_at(row, columns["material"])),
            text(value_at(row, columns["material_type"])),
            text(value_at(row, columns["weight"])),
            text(value_at(row, columns["cavity"])),
            text(value_at(row, columns["sets"])),
            machine_code,
            combined_capacity_machine or text(value_at(row, columns["target"])),
            text(value_at(row, columns["mold_base_type"])),
            text(value_at(row, columns["mold_base_material"])),
            text(value_at(row, columns["structure"])),
            text(value_at(row, columns["process"])),
            text(value_at(row, columns["cycle"])),
            text(value_at(row, columns["mold_size"])),
            text(value_at(row, columns["mold_specification"])),
            text(value_at(row, columns["color"])),
        )
        # 空白行不能沿用上一行模号；Water Table 模价表在 M12 后还有
        # 多个空白/说明行，旧逻辑会把这些行全部错误解析成 M12。
        if not raw_mold_no and price is None and not any(continuation_values):
            continue
        mold_no = raw_mold_no or last_mold_no
        if raw_mold_no:
            last_mold_no = raw_mold_no
        if not mold_no and not name and not chinese_name:
            continue
        if not name:
            name = chinese_name or mold_no
        if price is None:
            warnings.append(f"第 {source_row} 行 {name} 未识别模具价格，已按 0 预览")
            price = Decimal("0")
        sets = number(value_at(row, columns["sets"]), Decimal("1")) or Decimal("1")
        quantity = sets
        if sets > 100:
            amortization_candidates.append(sets)
            quantity = Decimal("1")
            warnings.append(f"第 {source_row} 行套数 {decimal_text(sets)} 较大，按分摊数量候选处理，模具数量按 1")
        row_images = (embedded_images_by_row or {}).get(source_row, [])
        image_reference = text(value_at(row, columns["image"]))
        if row_images:
            image_reference = "；".join(image.file_name for image in row_images)
        molds.append(
            {
                "item": name,
                "mold_no": mold_no,
                "chinese_name": chinese_name,
                "quantity": decimal_text(max(quantity, Decimal("0"))),
                "cost_rmb": decimal_text(max(price, Decimal("0"))),
                "material": text(value_at(row, columns["material"])),
                "material_type": text(value_at(row, columns["material_type"])),
                "net_weight_g": decimal_text(number(value_at(row, columns["weight"]), Decimal("0"))),
                "cavity": text(value_at(row, columns["cavity"])),
                "machine_code": machine_code,
                "target_output": decimal_text(target_output),
                "mold_base_type": text(value_at(row, columns["mold_base_type"])),
                "mold_base_material": text(value_at(row, columns["mold_base_material"])),
                "structure": text(value_at(row, columns["structure"])),
                "process": text(value_at(row, columns["process"])),
                "cycle_time_seconds": decimal_text(number(value_at(row, columns["cycle"]), Decimal("0"))),
                "mold_size": text(value_at(row, columns["mold_size"])),
                "mold_specification": text(value_at(row, columns["mold_specification"])),
                "color": text(value_at(row, columns["color"])),
                "image_reference": image_reference,
                "image_attachment_ids": [],
                "remark": text(value_at(row, columns["note"])),
                "source_row": source_row,
            }
        )
    if not molds:
        if is_zhanxing_mold_layout:
            amortization_qty = max(fallback_qty, Decimal("1"))
            warnings.append("展兴工模报价模板尚未填写模具明细")
            return {
                "molds": [],
                "amortization_qty": decimal_text(amortization_qty),
            }, 0, warnings
        raise ValueError("已识别模具表头，但没有解析到模具明细")

    amortization_qty = explicit_amortization
    if amortization_qty is None and amortization_candidates:
        amortization_qty = amortization_candidates[0]
    if amortization_qty is None:
        amortization_qty = max(fallback_qty, Decimal("1"))
        warnings.append(f"未识别模具分摊数量，预览按报价数量 {decimal_text(amortization_qty)} 填入")
    embedded_count = sum(len(items) for items in (embedded_images_by_row or {}).values())
    image_column_label = (
        get_column_letter(columns["image"] + 1)
        if columns["image"] is not None
        else ""
    )
    if embedded_count:
        warnings.append(
            f"已识别并提取 {image_column_label} 列 {embedded_count} 张嵌入图片；"
            "确认导入后将自动保存到对应模具行"
        )
    elif is_water_table_layout:
        warnings.append("U 列未识别到可提取的嵌入图片；如源表图片为链接，请改为嵌入图片后重试")
    elif is_zhanxing_mold_layout:
        warnings.append("F 列未识别到可提取的嵌入图片；如源表图片为链接，请改为嵌入图片后重试")
    else:
        warnings.append("嵌入模具图片不自动写入报价，请在对应工程分段上传图片附件")
    return {"molds": molds, "amortization_qty": decimal_text(amortization_qty)}, len(molds), warnings


def _parse_hardware(
    rows: list[list[object]],
    header_index: int,
    *,
    rmb_hkd: Decimal,
    **_: object,
) -> tuple[dict[str, Any], int, list[str]]:
    header = rows[header_index]
    columns = {
        "name": column_index(header, ("零件名称", "配件名称", "名称")),
        "specification": column_index(header, ("规格", "型号")),
        "quantity": column_index(header, ("用量", "数量")),
        "unit_price": column_index(header, ("单价RMB", "单价人民币", "人民币单价", "单价")),
        "remark": column_index(header, ("备注", "说明")),
    }
    unit_header = normalized(value_at(header, columns["unit_price"]))
    source_is_hkd = "HKD" in unit_header or "港币" in unit_header
    if source_is_hkd and rmb_hkd <= 0:
        raise ValueError("报价参考快照的 RMB/HKD 汇率无效")

    output: list[dict[str, Any]] = []
    warnings: list[str] = []
    for source_row, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        name = text(value_at(row, columns["name"]))
        specification = text(value_at(row, columns["specification"]))
        row_label = "".join(text(cell) for cell in row)
        if re.search(r"^(合计|小计|总计|说明|备注|附[:：])", row_label):
            continue
        if not name and not specification:
            continue
        quantity = number(value_at(row, columns["quantity"]))
        unit_price = number(value_at(row, columns["unit_price"]))
        if quantity is None:
            quantity = Decimal("0")
            warnings.append(f"第 {source_row} 行 {name or specification} 未识别用量，已按 0 预览")
        if unit_price is None:
            unit_price = Decimal("0")
            warnings.append(f"第 {source_row} 行 {name or specification} 未识别人民币单价，已按 0 预览")
        if quantity < 0 or unit_price < 0:
            raise ValueError(f"第 {source_row} 行 {name or specification} 的用量或单价不能小于 0")
        unit_price_rmb = unit_price * rmb_hkd if source_is_hkd else unit_price
        output.append(
            {
                "item": name,
                "category": "hardware",
                "specification": specification,
                "quantity": decimal_text(quantity),
                "unit_price_rmb": decimal_text(unit_price_rmb),
                "auxiliary_category": "五金",
                "tax_rate_percent": "13.0000",
                "remark": text(value_at(row, columns["remark"])),
                "source_row": source_row,
            }
        )
    if not output:
        raise ValueError("已识别五金表头，但没有解析到五金明细")
    if source_is_hkd:
        warnings.append(f"五金表使用 HKD 单价，已按快照 RMB/HKD={decimal_text(rmb_hkd)} 反算为人民币单价")
    warnings.append("五金导入仅映射零件名称、规格、用量、单价 RMB 和备注；类别固定为五金、税点固定为 13%，模板其他列不写入报价字段")
    warnings.append("模板内嵌图片不自动写入明细；如需保留，请在工程部分段上传原表或图片附件")
    return {"materials": output}, len(output), warnings


def _parse_electronic(
    rows: list[list[object]],
    header_index: int,
    *,
    rmb_hkd: Decimal,
    **_: object,
) -> tuple[dict[str, Any], int, list[str]]:
    header = rows[header_index]
    columns = {
        "name": column_index(header, ("零件名称",)),
        "spec": column_index(header, ("规格",)),
        "qty": column_index(header, ("用量", "数量")),
        "unit": column_index(header, ("单价",)),
        "amount": column_index(header, ("合计", "金额")),
        "tax": column_index(header, ("税点", "税率")),
        "note": column_index(header, ("备注",)),
    }
    unit_header = normalized(value_at(header, columns["unit"]))
    source_is_hkd = "HKD" in unit_header or "港币" in unit_header
    divisor = Decimal("1") if source_is_hkd else rmb_hkd
    if divisor <= 0:
        raise ValueError("报价参考快照的 RMB/HKD 汇率无效")

    scalar_labels = {
        "邦定": ("bonding_rmb", "bonding_hkd"),
        "邦定成本": ("bonding_rmb", "bonding_hkd"),
        "贴片": ("smt_rmb", "smt_hkd"),
        "贴片成本": ("smt_rmb", "smt_hkd"),
        "SMT": ("smt_rmb", "smt_hkd"),
        "SMT成本": ("smt_rmb", "smt_hkd"),
        "人工成本": ("labor_rmb", "labor_hkd"),
        "测试": ("testing_rmb", "testing_hkd"),
        "测试费用": ("testing_rmb", "testing_hkd"),
        "包装运输": ("packaging_rmb", "packaging_hkd"),
        "抵税差额": ("tax_credit_difference_rmb", "tax_credit_difference_hkd"),
    }
    fragment: dict[str, Any] = {"pricing_currency": "RMB", "components": []}
    for row in rows:
        row_text = "|".join(text(cell) for cell in row)
        for index, cell in enumerate(row):
            targets = scalar_labels.get(normalized(cell).rstrip(":："))
            if targets is not None and targets[0] not in fragment:
                parsed = _next_number(row, index)
                if parsed is not None:
                    amount_rmb = parsed * rmb_hkd if source_is_hkd else parsed
                    amount_hkd = parsed if source_is_hkd else parsed / divisor
                    fragment[targets[0]] = decimal_text(amount_rmb)
                    fragment[targets[1]] = decimal_text(amount_hkd)
        profit = re.search(r"(\d{1,3}(?:\.\d+)?)%\s*利润", row_text)
        if profit:
            fragment["profit_rate_percent"] = decimal_text(Decimal(profit.group(1)))

    warnings: list[str] = []
    last_parent: dict[str, Any] | None = None
    parsed_count = 0
    stop_labels = set(scalar_labels) | {
        "零件成本", "含税报价", "合计成本", "成本合计", "含利润价", "应交税负",
        "报价人", "审核", "核准",
    }
    for source_row, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        name = text(value_at(row, columns["name"]))
        spec = text(value_at(row, columns["spec"]))
        qty = number(value_at(row, columns["qty"]))
        unit = number(value_at(row, columns["unit"]))
        amount_value = number(value_at(row, columns["amount"]))
        if any(normalized(cell).rstrip(":：") in stop_labels for cell in row if text(cell)):
            continue
        if unit is None and amount_value is not None and qty and qty > 0:
            unit = amount_value / qty
        if not name and not spec:
            continue
        if qty is None and unit is None and amount_value is None:
            continue
        if qty is None:
            qty = Decimal("1")
        if unit is None:
            unit = Decimal("0")
            warnings.append(f"第 {source_row} 行 {name or spec} 未识别单价，已按 0 预览")
        component = {
            "item": name,
            "specification": spec,
            "quantity": decimal_text(max(qty, Decimal("0"))),
            "unit_price_rmb": decimal_text(max(unit * rmb_hkd if source_is_hkd else unit, Decimal("0"))),
            "unit_price_hkd": decimal_text(max(unit / divisor, Decimal("0"))),
            "tax_rate_percent": decimal_text(number(value_at(row, columns["tax"]), Decimal("13"))),
            "source_unit_price": decimal_text(unit),
            "source_currency": "HKD" if source_is_hkd else "RMB",
            "remark": text(value_at(row, columns["note"])),
            "source_row": source_row,
            "children": [],
        }
        parsed_count += 1
        if not name and spec and last_parent is not None:
            last_parent["children"].append(component)
            continue
        if name:
            fragment["components"].append(component)
            last_parent = component
    if not fragment["components"]:
        raise ValueError("已识别电子表头，但没有解析到电子零件明细")
    if not source_is_hkd:
        warnings.append(f"电子表人民币单价已原值导入，并按快照 RMB/HKD={decimal_text(rmb_hkd)} 自动换算港币预览")
    else:
        warnings.append(f"电子表港币单价已按快照 RMB/HKD={decimal_text(rmb_hkd)} 反算人民币单价")
    warnings.append("税点未填写的电子零件按 13% 导入；抵税差额、税负及含税报价由服务端按模板公式重算")
    return fragment, parsed_count, warnings


def _molding_header_kind(row: list[object]) -> str:
    if row_has(row, "模具名称|模号", "啤净重|净重", "材质|材料"):
        return "injection"
    if row_has(row, "货名", "日产量|预估料重", "吹工|披锋|利润"):
        return "blow"
    return ""


def _split_material_grade(material: str, grade: str) -> tuple[str, str]:
    if grade or not material:
        return material, grade
    parts = material.split()
    if len(parts) >= 2:
        return parts[0], " ".join(parts[1:])
    return material, grade


def _molding_loss_rate(header: list[object]) -> Decimal:
    for value in header:
        match = re.search(r"料损耗\s*(\d+(?:\.\d+)?)\s*%", text(value))
        if match:
            return Decimal(match.group(1))
    return Decimal("3")


def _parse_molding(
    rows: list[list[object]],
    header_index: int,
    **_: object,
) -> tuple[dict[str, Any], int, list[str]]:
    headers = [
        (index, _molding_header_kind(row))
        for index, row in enumerate(rows)
        if _molding_header_kind(row)
    ]
    if not headers:
        raise ValueError("未找到可识别的注塑或吹气报价表头")

    injection_lines: list[dict[str, Any]] = []
    blow_lines: list[dict[str, Any]] = []
    warnings: list[str] = []
    section_loss_rate = Decimal("3")
    for header_position, (section_header_index, kind) in enumerate(headers):
        header = rows[section_header_index]
        end_index = headers[header_position + 1][0] if header_position + 1 < len(headers) else len(rows)
        if kind == "injection":
            section_loss_rate = _molding_loss_rate(header)
            columns = {
                "item": preferred_column_index(header, ("模具名称", "项目", "零件名称")),
                "mold_no": preferred_column_index(header, ("模号", "模具编号", "MOLD NO")),
                "material": preferred_column_index(header, ("材质", "材料")),
                "grade": preferred_column_index(header, ("料型", "牌号", "材料型号")),
                "color": preferred_column_index(header, ("颜色",)),
                "weight": preferred_column_index(header, ("啤净重(G)", "啤净重", "净重(G)", "净重")),
                "machine_name": preferred_column_index(header, ("机台", "普通机")),
                "machine_code": preferred_column_index(header, ("机型A码", "A码", "机型")),
                "cavity": preferred_column_index(header, ("出模数", "穴数", "CAVITY")),
                "sets": preferred_column_index(header, ("套数",)),
                "target": preferred_column_index(header, ("目标数", "目标产量")),
                "cycle": preferred_column_index(header, ("周期(秒)", "周期", "CYCLE TIME")),
                "quantity": preferred_column_index(header, ("成品用量", "数量", "用量")),
                "remark": preferred_column_index(header, ("备注", "说明")),
            }
            for source_row, row in enumerate(rows[section_header_index + 1 : end_index], start=section_header_index + 2):
                item = text(value_at(row, columns["item"]))
                if not item or re.search(r"合计|汇总|参考表|料价表|机型价表", item):
                    continue
                material = text(value_at(row, columns["material"]))
                grade = text(value_at(row, columns["grade"]))
                material, grade = _split_material_grade(material, grade)
                weight = number(value_at(row, columns["weight"]))
                machine_code = text(value_at(row, columns["machine_code"]))
                target = number(value_at(row, columns["target"]), Decimal("0")) or Decimal("0")
                if not material and weight is None and not machine_code:
                    continue
                if not grade:
                    warnings.append(f"第 {source_row} 行 {item} 未识别料型；保存后需补齐材质 + 料型才能匹配冻结材料价")
                if not machine_code:
                    warnings.append(f"第 {source_row} 行 {item} 未识别机型 A 码")
                if target <= 0:
                    warnings.append(f"第 {source_row} 行 {item} 未识别目标数；保存前必须填写大于 0 的目标数")
                injection_lines.append({
                    "item": item,
                    "mold_no": text(value_at(row, columns["mold_no"])),
                    "material": material,
                    "grade": grade,
                    "color": text(value_at(row, columns["color"])),
                    "net_weight_g": decimal_text(max(weight or Decimal("0"), Decimal("0"))),
                    "loss_rate_percent": decimal_text(section_loss_rate),
                    "machine_name": text(value_at(row, columns["machine_name"])),
                    "machine_code": machine_code,
                    "cavity": text(value_at(row, columns["cavity"])),
                    "sets": decimal_text(max(number(value_at(row, columns["sets"]), Decimal("1")) or Decimal("1"), Decimal("0"))),
                    "target_output": decimal_text(max(target, Decimal("0"))),
                    "cycle_time_seconds": decimal_text(max(number(value_at(row, columns["cycle"]), Decimal("0")) or Decimal("0"), Decimal("0"))),
                    "quantity": decimal_text(max(number(value_at(row, columns["quantity"]), Decimal("1")) or Decimal("1"), Decimal("0"))),
                    "remark": text(value_at(row, columns["remark"])),
                    "source_row": source_row,
                })
        else:
            columns = {
                "item": preferred_column_index(header, ("货名", "项目", "产品名称")),
                "daily_capacity": preferred_column_index(header, ("日产量/22H", "日产量", "产能")),
                "material": preferred_column_index(header, ("用料", "材料", "材质")),
                "grade": preferred_column_index(header, ("料型", "牌号", "材料型号")),
                "weight": preferred_column_index(header, ("预估料重G", "预估料重", "料重G", "料重")),
                "labor": preferred_column_index(header, ("吹工", "人工HKD", "吹气人工")),
                "burr": preferred_column_index(header, ("披锋", "披锋HKD")),
                "profit": preferred_column_index(header, ("利润×", "利润倍率", "利润")),
                "quantity": preferred_column_index(header, ("成品用量", "数量")),
                "output_count": preferred_column_index(header, ("出数", "出模数")),
                "mold_price": preferred_column_index(header, ("模价(¥)", "模价RMB", "模价")),
                "remark": preferred_column_index(header, ("备注", "说明")),
            }
            for source_row, row in enumerate(rows[section_header_index + 1 : end_index], start=section_header_index + 2):
                item = text(value_at(row, columns["item"]))
                if not item or re.search(r"合计|汇总|参考表", item):
                    continue
                material = text(value_at(row, columns["material"]))
                grade = text(value_at(row, columns["grade"]))
                material, grade = _split_material_grade(material, grade)
                weight = number(value_at(row, columns["weight"]))
                if not material and weight is None:
                    continue
                if not grade:
                    warnings.append(f"第 {source_row} 行 {item} 未识别料型；保存后需补齐材质 + 料型才能匹配冻结材料价")
                blow_lines.append({
                    "item": item,
                    "daily_capacity": text(value_at(row, columns["daily_capacity"])),
                    "material": material,
                    "grade": grade,
                    "estimated_weight_g": decimal_text(max(weight or Decimal("0"), Decimal("0"))),
                    "labor_hkd": decimal_text(max(number(value_at(row, columns["labor"]), Decimal("0")) or Decimal("0"), Decimal("0"))),
                    "burr_hkd": decimal_text(max(number(value_at(row, columns["burr"]), Decimal("0")) or Decimal("0"), Decimal("0"))),
                    "profit_multiplier": decimal_text(max(number(value_at(row, columns["profit"]), Decimal("1.05")) or Decimal("1.05"), Decimal("0"))),
                    "quantity": decimal_text(max(number(value_at(row, columns["quantity"]), Decimal("1")) or Decimal("1"), Decimal("0"))),
                    "output_count": text(value_at(row, columns["output_count"])),
                    "mold_price_rmb": decimal_text(max(number(value_at(row, columns["mold_price"]), Decimal("0")) or Decimal("0"), Decimal("0"))),
                    "remark": text(value_at(row, columns["remark"])),
                    "source_row": source_row,
                })

    if not injection_lines and not blow_lines:
        raise ValueError("已识别啤机报价表头，但没有解析到注塑或吹气明细")
    warnings.append("导入表内料价、原料单价、啤价、产品料价、小计和合计不直接写入；保存后统一按本报价冻结材料价、机型价与 rr2-2026-v1 公式重算")
    if not injection_lines:
        warnings.append("本次未识别注塑明细；替换导入会清空现有注塑明细")
    if not blow_lines:
        warnings.append("本次未识别吹气明细；替换导入会清空现有吹气明细")
    fragment = {
        "injection_loss_rate_percent": decimal_text(section_loss_rate),
        "injection_lines": injection_lines,
        "blow_lines": blow_lines,
    }
    return fragment, len(injection_lines) + len(blow_lines), warnings


def _parse_painting(
    rows: list[list[object]],
    header_index: int,
    **_: object,
) -> tuple[dict[str, Any], int, list[str]]:
    header = rows[header_index]
    image_column = preferred_column_index(header, ("图片", "图", "图片引用", "附件引用"))
    name_column = preferred_column_index(header, ("名称", "零件名称", "产品名称"))
    position_column = preferred_column_index(header, ("位置", "部位"))
    note_column = preferred_column_index(header, ("备注", "说明"))
    process_columns: list[tuple[str, str, int, int]] = []
    for label, key in PAINTING_PROCESSES:
        qty_column = next(
            (
                index
                for index, value in enumerate(header)
                if normalized(label) in normalized(value) and "单价" not in normalized(value)
            ),
            None,
        )
        if qty_column is not None:
            unit_column = next(
                (
                    index
                    for index, value in enumerate(header)
                    if normalized(label) in normalized(value)
                    and any(token in normalized(value) for token in ("单价", "价格"))
                ),
                qty_column + 1,
            )
            process_columns.append((label, key, qty_column, unit_column))
    output: list[dict[str, Any]] = []
    warnings: list[str] = []
    for source_row, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        name = text(value_at(row, name_column))
        position = text(value_at(row, position_column))
        row_label = position or name
        if not row_label or re.search(r"合计|小计|总报价|[:：]$", row_label):
            continue
        operations = {
            key: {"quantity": "0.0000", "unit_price_hkd": "0.0000"}
            for _, key in PAINTING_PROCESSES
        }
        has_quantity = False
        for label, key, qty_column, unit_column in process_columns:
            qty = number(value_at(row, qty_column), Decimal("0")) or Decimal("0")
            unit = number(value_at(row, unit_column), Decimal("0")) or Decimal("0")
            if qty > 0:
                has_quantity = True
                if unit <= 0:
                    warnings.append(f"第 {source_row} 行 {row_label}/{label} 未识别单价，已按 0 预览")
            operations[key] = {
                "quantity": decimal_text(max(qty, Decimal("0"))),
                "unit_price_hkd": decimal_text(max(unit, Decimal("0"))),
            }
        if has_quantity:
            output.append(
                {
                    "image_reference": text(value_at(row, image_column)),
                    "name": name,
                    "position": position,
                    "operations": operations,
                    "remark": text(value_at(row, note_column)),
                    "source_row": source_row,
                }
            )
    if not output:
        raise ValueError("已识别喷油表头，但没有解析到喷油工序明细")
    warnings.append("源表总报价和合计仅用于核对，不直接导入；保存后按含 UV 的九类工序数量 × 单价由服务端重算")
    warnings.append("嵌入喷油图片不自动写入报价；图片单元格文本会保存为附件引用，原报价单可另存为分段附件")
    return {"rows": output}, len(output), warnings


def _parse_slush(
    rows: list[list[object]],
    header_index: int,
    **_: object,
) -> tuple[dict[str, Any], int, list[str]]:
    header = rows[header_index]
    columns = {
        "product_code": preferred_column_index(header, ("产品编号", "产品编码", "产品号", "货号")),
        "item": preferred_column_index(header, ("胶件名称", "零件名称", "产品名称", "名称")),
        "material": preferred_column_index(header, ("材料", "材质")),
        "weight_g": preferred_column_index(header, ("料重(G)", "料重", "净重(G)", "净重")),
        "daily_output_24h": preferred_column_index(header, ("日产量24H", "日产量", "24H产量")),
        "quantity": preferred_column_index(header, ("用量(PC)", "用量", "数量")),
        "unit_price_hkd": preferred_column_index(header, ("单价HKD", "单价HKS", "单价HK$", "单价")),
        "remark": preferred_column_index(header, ("备注", "说明")),
    }
    output: list[dict[str, Any]] = []
    warnings: list[str] = []
    for source_row, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        product_code = text(value_at(row, columns["product_code"]))
        item = text(value_at(row, columns["item"]))
        label = f"{product_code} {item}".strip()
        if not label or re.search(r"合计|小计|总价|总计", label):
            continue
        quantity = number(value_at(row, columns["quantity"]))
        unit_price = number(value_at(row, columns["unit_price_hkd"]))
        if quantity is None:
            quantity = Decimal("0")
            warnings.append(f"第 {source_row} 行 {label} 未识别用量，已按 0 预览")
        if unit_price is None:
            unit_price = Decimal("0")
            warnings.append(f"第 {source_row} 行 {label} 未识别单价 HKD，已按 0 预览")
        output.append(
            {
                "product_code": product_code,
                "item": item,
                "material": text(value_at(row, columns["material"])),
                "weight_g": decimal_text(max(number(value_at(row, columns["weight_g"]), Decimal("0")) or Decimal("0"), Decimal("0"))),
                "daily_output_24h": decimal_text(max(number(value_at(row, columns["daily_output_24h"]), Decimal("0")) or Decimal("0"), Decimal("0"))),
                "quantity": decimal_text(max(quantity, Decimal("0"))),
                "unit_price_hkd": decimal_text(max(unit_price, Decimal("0"))),
                "remark": text(value_at(row, columns["remark"])),
                "source_row": source_row,
            }
        )
    if not output:
        raise ValueError("已识别搪胶表头，但没有解析到搪胶产品明细")
    warnings.append("源表总价和合计仅用于核对，不直接导入；保存后按用量 × 单价 HKD 由服务端重算")
    return {"lines": output}, len(output), warnings


def _parse_sewing(
    rows: list[list[object]],
    header_index: int,
    **_: object,
) -> tuple[dict[str, Any], int, list[str]]:
    groups: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    columns: dict[str, int | None] | None = None
    pending_title = ""
    last_material = ""
    next_detail_row_is_product_title = False
    warnings: list[str] = []
    total_rows = 0

    for source_index, row in enumerate(rows, start=1):
        nonempty = [text(value) for value in row if text(value)]
        if row_has(row, "物料名称|布料名称", "裁片部位|部位", "用量|用量/码", "价钱|总价钱"):
            total_price_column = preferred_column_index(row, ("总价钱(RMB)", "总价钱"))
            if total_price_column is None:
                total_price_column = preferred_column_index(row, ("价钱(RMB)", "价钱", "成本"))
            columns = {
                "material": preferred_column_index(row, ("布料名称", "物料名称")),
                "part": preferred_column_index(row, ("裁片部位", "部位")),
                "craft": preferred_column_index(row, ("工艺",)),
                "pieces": preferred_column_index(row, ("裁片数",)),
                "supplier": column_index(row, ("供应商",)),
                "fabric_moq_y": preferred_column_index(row, ("布料MOQ/Y", "MOQ/Y", "布料MOQ")),
                "below_moq_fee_rmb": preferred_column_index(
                    row,
                    ("低于MOQ/每色产生费用", "低于MOQ每色产生费用", "低于MOQ费用"),
                ),
                "usage": column_index(row, ("用量",)),
                "unit": preferred_column_index(row, ("物料价(RMB)", "物料价", "单价(RMB)", "单价")),
                "exchange_rate": preferred_column_index(row, ("汇率", "RMB/HKD汇率", "RMB→HKD汇率")),
                "markup": column_index(row, ("码点",)),
                "price": total_price_column,
                "note": column_index(row, ("备注",)),
            }
            if current is None or current.get("materials"):
                name = pending_title or f"导入产品 {len(groups) + 1}"
                current = {
                    "name": name,
                    "category": "hair" if "发" in name else "clothes",
                    "materials": [],
                    "labor_rmb": "0.0000",
                }
                groups.append(current)
            pending_title = ""
            last_material = ""
            next_detail_row_is_product_title = False
            continue
        if current is None and nonempty and len(nonempty) <= 2:
            candidate = " ".join(nonempty)
            if not re.search(r"^明细表|^DATE|日期|^20\d\d", candidate, re.I):
                pending_title = re.sub(r"^产品\s*[:：]\s*", "", candidate).strip()
            continue
        if current is None or columns is None:
            continue
        raw_material = text(value_at(row, columns["material"]))
        usage = number(value_at(row, columns["usage"]))
        unit = number(value_at(row, columns["unit"]))
        exchange_rate = number(value_at(row, columns["exchange_rate"]))
        price = number(value_at(row, columns["price"]))
        part = text(value_at(row, columns["part"]))
        if any("合计" in value for value in nonempty):
            next_detail_row_is_product_title = True
            last_material = ""
            continue
        if next_detail_row_is_product_title and raw_material and usage is None and unit is None and price is None and not part:
            product_name = re.sub(r"^产品\s*[:：]\s*", "", raw_material).strip()
            current = {
                "name": product_name,
                "category": "hair" if "发" in product_name else "clothes",
                "materials": [],
                "labor_rmb": "0.0000",
            }
            groups.append(current)
            next_detail_row_is_product_title = False
            last_material = ""
            continue
        if raw_material and usage is None and unit is None and price is None and not part:
            if not current["materials"]:
                current["name"] = raw_material
                current["category"] = "hair" if "发" in raw_material else "clothes"
                last_material = ""
                continue
        if not nonempty:
            last_material = ""
            continue
        material = raw_material
        if material:
            last_material = material
        elif part and last_material:
            # Source sewing quotes commonly write a material only on the first
            # cutting-part row.  Following rows belong to the same material even
            # though column A is blank, so forward-fill it within this detail block.
            material = last_material
        if not material:
            continue
        usage_value = max(usage or Decimal("1"), Decimal("0"))
        markup_value = max(number(value_at(row, columns["markup"]), Decimal("1")) or Decimal("1"), Decimal("0"))
        if markup_value <= 0:
            markup_value = Decimal("1")
        if unit is None:
            price_to_rmb = exchange_rate if exchange_rate is not None and exchange_rate > 0 else Decimal("1")
            unit = (
                price * price_to_rmb / usage_value / markup_value
                if price is not None and usage_value > 0
                else Decimal("0")
            )
            warnings.append(f"第 {source_index} 行 {material} 未识别物料价，已按源表总价反算或按 0 预览")
        craft_text = text(value_at(row, columns["craft"]))
        current["materials"].append(
            {
                "item": material,
                "part": part,
                "craft": "电绣" if "电绣" in craft_text else "丝印" if "丝印" in craft_text else "",
                "pieces": decimal_text(max(number(value_at(row, columns["pieces"]), Decimal("0")) or Decimal("0"), Decimal("0"))),
                "supplier": text(value_at(row, columns["supplier"])),
                **(
                    {"fabric_moq_y": precise_decimal_text(
                        max(number(value_at(row, columns["fabric_moq_y"]), Decimal("0")) or Decimal("0"), Decimal("0"))
                    )}
                    if columns["fabric_moq_y"] is not None else {}
                ),
                **(
                    {"below_moq_fee_rmb": precise_decimal_text(
                        max(number(value_at(row, columns["below_moq_fee_rmb"]), Decimal("0")) or Decimal("0"), Decimal("0"))
                    )}
                    if columns["below_moq_fee_rmb"] is not None else {}
                ),
                "usage": precise_decimal_text(usage_value),
                "unit_price_rmb": precise_decimal_text(max(unit, Decimal("0"))),
                **(
                    {"exchange_rate": precise_decimal_text(max(exchange_rate or Decimal("0"), Decimal("0")))}
                    if columns["exchange_rate"] is not None and exchange_rate is not None and exchange_rate > 0 else {}
                ),
                "markup": precise_decimal_text(markup_value),
                "remark": text(value_at(row, columns["note"])),
                "source_row": source_index,
            }
        )
        total_rows += 1
    groups = [group for group in groups if group["materials"]]
    if not groups:
        raise ValueError("已识别车缝表头，但没有解析到车缝产品分组")
    warnings.append(
        "源表成本、总价钱和合计仅用于核对；保存后按用量/码 × 单价 RMB ÷ 行汇率 × 码点由服务端重算 HKD，裁片数不参与金额"
    )
    return {"groups": groups}, total_rows, warnings


def _assembly_pairs(header: list[object]) -> list[tuple[int, int, str]]:
    pairs: list[tuple[int, int, str]] = []
    pending: tuple[int, str] | None = None
    for index, value in enumerate(header):
        token = text(value)
        if "工序名称" in token:
            pending = (index, token.replace("工序名称", "").strip() or "装配")
        elif "人数" in token and pending:
            pairs.append((pending[0], index, pending[1]))
            pending = None
    return pairs


def _workshop_assembly_columns(header: list[object]) -> dict[str, int] | None:
    columns = {
        "group": preferred_column_index(header, ("货号或图片",)),
        "process": preferred_column_index(header, ("做工名称",)),
        "production_qty": preferred_column_index(header, ("总目标数量",)),
        "persons": preferred_column_index(header, ("人数",)),
        "remark": preferred_column_index(
            header,
            ("备注（看图、手办、样板、参考办）", "备注"),
        ),
    }
    required = ("group", "process", "production_qty", "persons")
    if any(columns[key] is None for key in required):
        return None
    return {
        key: int(value)
        for key, value in columns.items()
        if value is not None
    }


def _parse_workshop_assembly(
    rows: list[list[object]],
    header_index: int,
    *,
    columns: dict[str, int],
    fallback_qty: Decimal,
) -> tuple[dict[str, Any], int, list[str]]:
    raw_groups: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    packaging_started = False
    warnings: list[str] = []
    conflicting_qty_groups: set[str] = set()
    total_rows = 0

    for source_row, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        group_name = text(value_at(row, columns["group"]))
        process_name = text(value_at(row, columns["process"]))
        production_qty = number(value_at(row, columns["production_qty"]))
        persons = number(value_at(row, columns["persons"]))
        if not process_name:
            # 车间表以蓝灰合计行结束当前区域：D/E 为空，F 为本区人数合计。
            # 重置当前组，确保后续即使使用相同产品名称也仍按新区域建组。
            if not group_name and production_qty is None and persons is not None and persons > 0:
                current = None
            continue
        if persons is None or persons <= 0:
            continue
        if re.search(r"做工名称|工序名称|车间填写|不用填|合计|小计|总计", process_name):
            if re.search(r"合计|小计|总计", process_name):
                current = None
            continue

        if group_name:
            if "包装" in group_name:
                packaging_started = True
            category = "packaging" if packaging_started else "assembly"
            if (
                current is None
                or current["name"] != group_name
                or current["category"] != category
            ):
                current = {
                    "name": group_name,
                    "category": category,
                    "_production_qty": None,
                    "teams": "1.0000",
                    "processes": [],
                }
                raw_groups.append(current)
        if current is None:
            warnings.append(f"第 {source_row} 行 {process_name} 未找到 C 列所属区域，已跳过")
            continue

        if production_qty is not None and production_qty > 0:
            current_qty = current.get("_production_qty")
            if current_qty is None:
                current["_production_qty"] = production_qty
            elif production_qty != current_qty and current["name"] not in conflicting_qty_groups:
                warnings.append(
                    f"{current['name']} 区域 E 列出现多个每日生产数量，"
                    f"预览采用首个值 {decimal_text(current_qty)}"
                )
                conflicting_qty_groups.add(current["name"])

        current["processes"].append(
            {
                "name": process_name,
                "persons": decimal_text(persons),
                "remark": text(value_at(row, columns.get("remark"))),
                "source_row": source_row,
            }
        )
        total_rows += 1

    groups: list[dict[str, Any]] = []
    for group in raw_groups:
        if not group["processes"]:
            continue
        production_qty = group.pop("_production_qty", None)
        if production_qty is None:
            production_qty = max(fallback_qty, Decimal("1"))
            warnings.append(
                f"{group['name']} 未识别 E 列每日生产数量，"
                f"预览按报价数量 {decimal_text(production_qty)} 填入"
            )
        group["production_qty"] = decimal_text(production_qty)
        groups.append(group)

    if not groups:
        raise ValueError("已识别装工表头，但没有解析到装配或包装工序")
    warnings.append(
        "已按 C 列区域、D 列做工名称、E 列每日生产数量、F 列工艺人数导入；"
        "C 列首次出现“包装”后，该区域及后续区域归入包装部分"
    )
    return {"groups": groups}, total_rows, warnings


def _parse_assembly(
    rows: list[list[object]],
    header_index: int,
    *,
    fallback_qty: Decimal,
    **_: object,
) -> tuple[dict[str, Any], int, list[str]]:
    workshop_columns = _workshop_assembly_columns(rows[header_index])
    if workshop_columns is not None:
        return _parse_workshop_assembly(
            rows,
            header_index,
            columns=workshop_columns,
            fallback_qty=fallback_qty,
        )

    groups: list[dict[str, Any]] = []
    warnings: list[str] = []
    total_rows = 0
    for name_column, people_column, group_name in _assembly_pairs(rows[header_index]):
        production_qty: Decimal | None = None
        for row in rows[header_index + 1 :]:
            label = text(value_at(row, people_column - 1))
            capacity = number(value_at(row, people_column))
            if "产能" in label and capacity and capacity > 0:
                production_qty = capacity
                break
        if production_qty is None:
            production_qty = max(fallback_qty, Decimal("1"))
            warnings.append(f"{group_name} 未识别产能，预览按报价数量 {decimal_text(production_qty)} 填入")
        processes: list[dict[str, Any]] = []
        for source_row, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            name = text(value_at(row, name_column))
            people = number(value_at(row, people_column))
            if not name or people is None or people <= 0:
                continue
            if re.search(r"序号|工序名称|生产拉线|产品图片|客名|货号|目标数|说明|注意事项", name):
                continue
            processes.append(
                {
                    "name": name,
                    "persons": decimal_text(people),
                    "remark": text(value_at(row, people_column + 2)),
                    "source_row": source_row,
                }
            )
        if processes:
            category = "packaging" if re.search(r"包装|混装|装箱|入箱|彩盒|外箱|吸塑", group_name) else "assembly"
            groups.append({
                "name": group_name,
                "category": category,
                "production_qty": decimal_text(production_qty),
                "teams": "1.0000",
                "processes": processes,
            })
            total_rows += len(processes)
    if not groups:
        raise ValueError("已识别装配表头，但没有解析到生产排拉工序")
    return {"groups": groups}, total_rows, warnings


def _parse_hair(
    rows: list[list[object]], header_index: int, *, sheet_name: str = "明细",
    fallback_product_name: str = "", **_: Any,
) -> tuple[dict[str, Any], int, list[str]]:
    header = rows[header_index]
    name_col = column_index(header, ("货名", "名称"))
    item_col = column_index(header, ("货号", "编号"))
    unit_col = column_index(header, ("单位",))
    # The product price and weight precede its unit column (D/E before F in
    # the fixed template). Restrict the search so a missing left-hand header
    # cannot silently fall back to the right-hand raw-material breakdown.
    product_header = header[:unit_col] if unit_col is not None else []
    price_col = column_index(product_header, ("单价(HK$)", "单价(HKD)", "单价港币", "单价(港币)"))
    weight_col = column_index(product_header, ("重量(g)",))
    craft_col = column_index(header, ("工艺",))
    remark_col = column_index(header, ("备注",))
    if price_col is None or weight_col is None:
        raise ValueError("车发明细缺少港币单价或重量(g)列，不能使用右侧材料明细单价代替")
    lines = []
    warnings = []
    for source_row, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        name = text(value_at(row, name_col))
        item_no = text(value_at(row, item_col))
        raw_price = value_at(row, price_col)
        raw_weight = value_at(row, weight_col)
        if not any(text(value) for value in (name, item_no, raw_price, raw_weight)):
            continue
        if re.search(r"合计|小计|总计", name + item_no):
            continue
        unit_price = number(raw_price)
        weight = number(raw_weight)
        if unit_price is None or weight is None or unit_price <= 0 or weight <= 0:
            raise ValueError(f"明细第 {source_row} 行的单价和重量必须为正数；公式请先在 Excel 中计算并保存")
        if not name:
            name = fallback_product_name or item_no or "车发报价"
            warnings.append(f"明细第 {source_row} 行未填货名，暂用‘{name}’，请核对")
        lines.append({
            "name": name,
            "craft": text(value_at(row, craft_col)) or "车发",
            "weight_g": precise_decimal_text(weight),
            "unit_price_hkd": precise_decimal_text(unit_price),
            "unit": text(value_at(row, unit_col)) or "PCS",
            "remark": text(value_at(row, remark_col)),
            "item_no": item_no,
            "source_sheet": sheet_name,
            "source_row": source_row,
        })
    if not lines:
        raise ValueError("车发明细页没有可导入的单价和重量，请填写后保存")
    return {"lines": lines}, len(lines), warnings


PARSERS: dict[str, Callable[..., tuple[dict[str, Any], int, list[str]]]] = {
    "mold": _parse_mold,
    "hardware": _parse_hardware,
    "electronic": _parse_electronic,
    "molding": _parse_molding,
    "painting": _parse_painting,
    "slush": _parse_slush,
    "sewing": _parse_sewing,
    "hair": _parse_hair,
    "assembly": _parse_assembly,
}


def parse_internal_quote_workbook(
    content: bytes,
    import_type: str,
    *,
    rmb_hkd: Decimal = Decimal("0.85"),
    fallback_qty: Decimal = Decimal("1"),
    fallback_product_name: str = "",
) -> ParsedInternalQuoteImport:
    if import_type not in IMPORT_TYPE_DEPARTMENTS:
        raise ValueError("不支持的内部报价导入类型")
    sheets = workbook_rows(content, allow_xls=import_type == "hair")
    sheet_name, rows, header_index = find_header(sheets, import_type)
    rows = [list(row) for row in rows]
    if header_index + 1 < len(rows):
        next_row = rows[header_index + 1]
        header_words = (
            "名称",
            "编号",
            "规格",
            "用量",
            "数量",
            "单价",
            "金额",
            "材质",
            "人数",
            "位置",
            "备注",
            "NAME",
            "MATERIAL",
            "AMOUNT",
            "CHINESE",
            "DESCRIPTION",
            "TYPE",
            "DIM",
            "WEIGHT",
            "INSERT",
            "SLIDE",
            "TIME",
            "RATE",
            "GATE",
            "PICTURES",
            "REMARKS",
        )
        next_text = "|".join(normalized(cell) for cell in next_row)
        hits = sum(1 for word in header_words if normalized(word) in next_text)
        numeric_cells = sum(1 for cell in next_row if number(cell) is not None)
        if hits >= 2 and numeric_cells <= 1:
            width = max(len(rows[header_index]), len(next_row))
            rows[header_index] = [
                " ".join(
                    value
                    for value in (
                        text(value_at(rows[header_index], column)),
                        text(value_at(next_row, column)),
                    )
                    if value
                )
                for column in range(width)
            ]
            rows[header_index + 1] = []
    embedded_images: list[EmbeddedWorkbookImage] = []
    embedded_images_by_row: dict[int, list[EmbeddedWorkbookImage]] = {}
    if import_type == "mold":
        image_column = column_index(
            rows[header_index],
            ("图片", "图  片", "IMAGE", "PICTURES"),
        )
        embedded_images = workbook_embedded_images(
            content,
            sheet_name,
            source_column=(image_column + 1) if image_column is not None else 21,
            min_source_row=header_index + 2,
        )
        for image in embedded_images:
            embedded_images_by_row.setdefault(image.source_row, []).append(image)
    fragment, row_count, warnings = PARSERS[import_type](
        rows,
        header_index,
        rmb_hkd=rmb_hkd,
        fallback_qty=fallback_qty,
        sheet_name=sheet_name,
        embedded_images_by_row=embedded_images_by_row,
        **({"fallback_product_name": fallback_product_name} if import_type == "hair" else {}),
    )
    if import_type == "mold":
        parsed_rows = fragment.get("molds", [])
        used_source_rows = {
            int(row.get("source_row", 0) or 0)
            for row in parsed_rows
            if isinstance(row, dict)
        }
        embedded_images = [image for image in embedded_images if image.source_row in used_source_rows]
    return ParsedInternalQuoteImport(
        target_department=IMPORT_TYPE_DEPARTMENTS[import_type],
        sheet_name=sheet_name,
        header_row=header_index + 1,
        row_count=row_count,
        payload_fragment=fragment,
        warnings=warnings,
        embedded_images=embedded_images,
    )
