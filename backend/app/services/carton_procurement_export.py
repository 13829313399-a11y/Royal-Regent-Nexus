from __future__ import annotations

import json
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models.carton_procurement import CartonOrder, CartonOrderLine, CartonPurchaseOrderIssue


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FACTORY_NAMES = {
    "huakang-a": "华康 A 厂",
    "huakang-b": "华康 B 厂",
    "huakang-c": "华康 C 厂",
    "huakang-d": "华康 D 厂",
    "huadeng": "华登厂",
    "huaxing": "华兴厂",
}

PURCHASE_ORDER_TYPE_LABELS = {
    "INITIAL": "首次采购单",
    "APPEND": "追加采购单",
    "REDUCE": "减单通知",
    "ADJUSTMENT": "采购变更单",
}


def build_purchase_order_issue_workbook(issue: CartonPurchaseOrderIssue) -> bytes:
    snapshot = json.loads(issue.snapshot_json)
    order = snapshot["order"]
    lines = snapshot["lines"]
    document_label = PURCHASE_ORDER_TYPE_LABELS.get(issue.document_type, "采购变更单")

    workbook = Workbook()
    workbook.properties.creator = "Royal Regent Nexus"
    workbook.properties.title = f"{issue.document_no} {document_label}"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    sheet = workbook.active
    sheet.title = document_label
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A10"

    teal = "087F74"
    teal_dark = "075E57"
    teal_light = "E8F6F4"
    amber_light = "FEF3C7"
    red_light = "FEE2E2"
    slate = "475569"
    slate_light = "F1F5F9"
    white = "FFFFFF"
    thin = Side(style="thin", color="CBD5E1")
    medium = Side(style="medium", color=teal)

    sheet.merge_cells("A1:J1")
    sheet["A1"] = f"Royal Regent Nexus · {FACTORY_NAMES.get(issue.factory_id, issue.factory_id)}{document_label}"
    sheet["A1"].font = Font(name="Microsoft YaHei", size=18, bold=True, color=white)
    sheet["A1"].fill = PatternFill("solid", fgColor=teal_dark)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 34

    sheet.merge_cells("A2:J2")
    sheet["A2"] = "PAPER PACKAGING PURCHASE ORDER CHANGE"
    sheet["A2"].font = Font(name="Arial", size=9, bold=True, color=teal_dark)
    sheet["A2"].fill = PatternFill("solid", fgColor=teal_light)
    sheet["A2"].alignment = Alignment(horizontal="center", vertical="center")

    information = {
        "A4": "采购单号",
        "B4": issue.document_no,
        "E4": "原合同订单号",
        "F4": issue.order_no,
        "I4": "单据类型",
        "J4": document_label,
        "A5": "供应商",
        "B5": order.get("supplier_name", ""),
        "E5": "客户 / 合同号",
        "F5": f"{order.get('customer_name', '')} / {order.get('contract_no', '')}",
        "I5": "货号",
        "J5": order.get("item_no", ""),
        "A6": "变更前产品数量",
        "B6": float(issue.before_product_quantity),
        "E6": "本次产品变化",
        "F6": float(issue.product_quantity_delta),
        "I6": "变更后累计数量",
        "J6": float(issue.after_product_quantity),
        "A7": "原计划交期",
        "B7": snapshot.get("before_due_date") or "—",
        "E7": "本次计划交期",
        "F7": snapshot.get("after_due_date") or "—",
        "I7": "产品名称",
        "J7": order.get("product_name") or "—",
    }
    for cell_ref, value in information.items():
        sheet[cell_ref] = value
    for merged_range in ("B4:D4", "F4:H4", "B5:D5", "F5:H5", "B6:D6", "F6:H6", "B7:D7", "F7:H7"):
        sheet.merge_cells(merged_range)
    label_cells = ("A4", "E4", "I4", "A5", "E5", "I5", "A6", "E6", "I6", "A7", "E7", "I7")
    for row in sheet.iter_rows(min_row=4, max_row=7, min_col=1, max_col=10):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.font = Font(name="Microsoft YaHei", size=9, color="0F172A")
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for cell_ref in label_cells:
        sheet[cell_ref].font = Font(name="Microsoft YaHei", size=9, bold=True, color=slate)
        sheet[cell_ref].fill = PatternFill("solid", fgColor=slate_light)
    for cell_ref in ("B6", "F6", "J6"):
        sheet[cell_ref].number_format = "+#,##0.######;-#,##0.######;0" if cell_ref == "F6" else "#,##0.######"

    headers = [
        "序号", "纸品类型", "纸质", "规格", "每箱个数",
        "变更前箱数", "本次变化", "变更后累计", "单位", "说明",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=9, column=column, value=header)
        cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=teal)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=medium, bottom=medium)

    first_line_row = 10
    for index, line in enumerate(lines, start=1):
        row = first_line_row + index - 1
        before_required = float(line["before_required_quantity"])
        delta = float(line["required_quantity_delta"])
        after_required = float(line["after_required_quantity"])
        values = [
            index,
            line["packaging_type"],
            line["paper_quality"],
            f"{line['specification']} {line.get('dimension_unit', '')}".strip(),
            float(line["usage_quantity"]),
            before_required,
            delta,
            after_required,
            line["unit"],
            "现有余量覆盖，无需新增" if issue.document_type == "APPEND" and delta == 0 else "",
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.font = Font(name="Microsoft YaHei", size=9, color="0F172A")
            cell.alignment = Alignment(
                horizontal="right" if column in {5, 6, 7, 8} else "center" if column in {1, 9} else "left",
                vertical="center",
                wrap_text=column in {4, 10},
            )
            cell.border = Border(left=thin, right=thin, bottom=thin)
        sheet.cell(row=row, column=7).number_format = "+#,##0.######;-#,##0.######;0"
        sheet.cell(row=row, column=7).fill = PatternFill(
            "solid", fgColor=red_light if delta < 0 else amber_light if delta > 0 else slate_light
        )
        sheet.cell(row=row, column=7).font = Font(name="Microsoft YaHei", size=10, bold=True, color="991B1B" if delta < 0 else "92400E")
        sheet.row_dimensions[row].height = 27

    notice_row = first_line_row + len(lines) + 1
    sheet.merge_cells(start_row=notice_row, start_column=1, end_row=notice_row, end_column=10)
    if issue.document_type == "INITIAL":
        notice = "供应商请按“本次变化”列执行首次采购；“变更后累计”用于核对。"
    elif issue.document_type == "REDUCE":
        notice = "供应商请按“本次变化”列核减未交数量；负数表示减少，累计列仅用于核对。"
    else:
        notice = "供应商请只按“本次变化”列增加或调整数量；不得把“变更后累计”重复作为新增订单。"
    sheet.cell(row=notice_row, column=1, value=notice)
    sheet.cell(row=notice_row, column=1).font = Font(name="Microsoft YaHei", size=10, bold=True, color=teal_dark)
    sheet.cell(row=notice_row, column=1).fill = PatternFill("solid", fgColor=teal_light)
    sheet.cell(row=notice_row, column=1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[notice_row].height = 34

    generated_row = notice_row + 2
    sheet.merge_cells(start_row=generated_row, start_column=1, end_row=generated_row, end_column=4)
    sheet.cell(row=generated_row, column=1, value=f"制单人：{issue.generated_by_name or '—'}")
    sheet.merge_cells(start_row=generated_row, start_column=5, end_row=generated_row, end_column=10)
    sheet.cell(row=generated_row, column=5, value=f"固定生成时间：{issue.generated_at[:16].replace('T', ' ')}")
    for cell_ref in (f"A{generated_row}", f"E{generated_row}"):
        sheet[cell_ref].font = Font(name="Microsoft YaHei", size=9, color=slate)

    widths = [8, 14, 16, 28, 13, 14, 14, 14, 10, 24]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.print_area = f"A1:J{generated_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.35
    sheet.page_margins.bottom = 0.35

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def build_purchase_order_issue_batch_workbook(
    issues: list[CartonPurchaseOrderIssue],
    *,
    generated_at: datetime,
) -> bytes:
    workbook = Workbook()
    workbook.properties.creator = "Royal Regent Nexus"
    workbook.properties.title = f"供应商采购单发行批次（{len(issues)} 份）"
    sheet = workbook.active
    sheet.title = "供应商采购单批次"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A6"

    teal = "087F74"
    teal_dark = "075E57"
    teal_light = "E8F6F4"
    white = "FFFFFF"
    thin = Side(style="thin", color="CBD5E1")
    sheet.merge_cells("A1:P1")
    sheet["A1"] = f"Royal Regent Nexus · {FACTORY_NAMES.get(issues[0].factory_id, issues[0].factory_id)}供应商采购单发行批次"
    sheet["A1"].font = Font(name="Microsoft YaHei", size=18, bold=True, color=white)
    sheet["A1"].fill = PatternFill("solid", fgColor=teal_dark)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 34
    sheet.merge_cells("A2:P2")
    sheet["A2"] = "每个采购单号均为独立固定快照；供应商只执行“本次箱数变化”，累计数量仅供核对。"
    sheet["A2"].font = Font(name="Microsoft YaHei", size=10, bold=True, color=teal_dark)
    sheet["A2"].fill = PatternFill("solid", fgColor=teal_light)
    sheet["A2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.merge_cells("A4:H4")
    sheet["A4"] = f"批次生成时间：{generated_at.strftime('%Y-%m-%d %H:%M')}"
    sheet.merge_cells("I4:P4")
    sheet["I4"] = f"独立采购单：{len(issues)} 份"

    headers = [
        "序号", "采购单号", "单据类型", "原合同订单号", "客户", "合同号", "货号", "计划交期",
        "纸品类型", "纸质", "规格", "每箱个数", "变更前箱数", "本次箱数变化", "变更后累计", "单位",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=5, column=column, value=header)
        cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=teal)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    current_row = 6
    for issue_index, issue in enumerate(issues, start=1):
        snapshot = json.loads(issue.snapshot_json)
        order = snapshot["order"]
        document_label = PURCHASE_ORDER_TYPE_LABELS.get(issue.document_type, "采购变更单")
        for line in snapshot["lines"]:
            values = [
                issue_index,
                issue.document_no,
                document_label,
                issue.order_no,
                order.get("customer_name", ""),
                order.get("contract_no", ""),
                order.get("item_no", ""),
                snapshot.get("after_due_date", ""),
                line["packaging_type"],
                line["paper_quality"],
                f"{line['specification']} {line.get('dimension_unit', '')}".strip(),
                float(line["usage_quantity"]),
                float(line["before_required_quantity"]),
                float(line["required_quantity_delta"]),
                float(line["after_required_quantity"]),
                line["unit"],
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=current_row, column=column, value=value)
                cell.font = Font(name="Microsoft YaHei", size=9, color="0F172A")
                cell.alignment = Alignment(
                    horizontal="right" if column in {12, 13, 14, 15} else "center" if column in {1, 3, 8, 16} else "left",
                    vertical="center",
                    wrap_text=column in {5, 6, 7, 11},
                )
                cell.border = Border(left=thin, right=thin, bottom=thin)
            sheet.cell(row=current_row, column=14).number_format = "+#,##0.######;-#,##0.######;0"
            sheet.cell(row=current_row, column=14).font = Font(name="Microsoft YaHei", size=10, bold=True, color="92400E")
            current_row += 1

    notice_row = current_row + 1
    sheet.merge_cells(start_row=notice_row, start_column=1, end_row=notice_row, end_column=16)
    sheet.cell(
        row=notice_row,
        column=1,
        value="严禁把“变更后累计”再次作为新增下单数量；首次、追加、减单均以采购单号和“本次箱数变化”列为执行依据。",
    )
    sheet.cell(row=notice_row, column=1).font = Font(name="Microsoft YaHei", size=10, bold=True, color=teal_dark)
    sheet.cell(row=notice_row, column=1).fill = PatternFill("solid", fgColor=teal_light)
    sheet.cell(row=notice_row, column=1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[notice_row].height = 34

    widths = [8, 24, 14, 22, 18, 20, 18, 13, 14, 16, 28, 13, 14, 16, 14, 10]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.auto_filter.ref = f"A5:P{current_row - 1}"
    sheet.print_area = f"A1:P{notice_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def build_purchase_order_workbook(
    order: CartonOrder,
    lines: list[CartonOrderLine],
    *,
    generated_at: datetime,
) -> bytes:
    workbook = Workbook()
    workbook.properties.creator = "Royal Regent Nexus"
    workbook.properties.title = f"{order.order_no} 纸箱采购单"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    sheet = workbook.active
    sheet.title = "纸箱采购单"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A10"

    teal = "087F74"
    teal_dark = "075E57"
    teal_light = "E8F6F4"
    slate = "475569"
    slate_light = "F1F5F9"
    border_color = "CBD5E1"
    white = "FFFFFF"
    thin = Side(style="thin", color=border_color)
    medium = Side(style="medium", color=teal)

    sheet.merge_cells("A1:H1")
    sheet["A1"] = f"Royal Regent Nexus · {FACTORY_NAMES.get(order.factory_id, order.factory_id)}纸箱采购单"
    sheet["A1"].font = Font(name="Microsoft YaHei", size=18, bold=True, color=white)
    sheet["A1"].fill = PatternFill("solid", fgColor=teal_dark)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 34

    sheet.merge_cells("A2:H2")
    sheet["A2"] = "PAPER PACKAGING PURCHASE ORDER"
    sheet["A2"].font = Font(name="Arial", size=9, bold=True, color=teal_dark)
    sheet["A2"].fill = PatternFill("solid", fgColor=teal_light)
    sheet["A2"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[2].height = 21

    information_rows = [
        (4, (("A4", "采购单号"), ("B4", order.order_no), ("D4", "下单日期"), ("E4", order.order_date), ("F4", "计划交期"), ("G4", order.due_date))),
        (5, (("A5", "采购厂区"), ("B5", FACTORY_NAMES.get(order.factory_id, order.factory_id)), ("D5", "供应商"), ("E5", order.supplier_name_snapshot))),
        (6, (("A6", "客户"), ("B6", order.customer_name), ("D6", "合同号"), ("E6", order.contract_no), ("G6", "货号"), ("H6", order.item_no))),
        (7, (("A7", "产品名称"), ("B7", order.product_name or "—"), ("D7", "产品订单数量"), ("E7", float(order.product_order_quantity)), ("G7", "流程状态"), ("H7", "已下单"))),
    ]
    for _, values in information_rows:
        for cell_ref, value in values:
            sheet[cell_ref] = value

    for merged_range in ("B4:C4", "G4:H4", "B5:C5", "E5:H5", "B6:C6", "E6:F6", "B7:C7", "E7:F7"):
        sheet.merge_cells(merged_range)

    label_cells = ("A4", "D4", "F4", "A5", "D5", "A6", "D6", "G6", "A7", "D7", "G7")
    for cell_ref in label_cells:
        cell = sheet[cell_ref]
        cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color=slate)
        cell.fill = PatternFill("solid", fgColor=slate_light)
        cell.alignment = Alignment(horizontal="left", vertical="center")
    for row in sheet.iter_rows(min_row=4, max_row=7, min_col=1, max_col=8):
        for cell in row:
            cell.border = Border(bottom=thin)
            if cell.coordinate not in label_cells:
                cell.font = Font(name="Microsoft YaHei", size=10, color="0F172A")
                cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet["E7"].number_format = (
        "#,##0" if order.product_order_quantity == order.product_order_quantity.to_integral_value() else "#,##0.######"
    )

    headers = ["序号", "纸品类型", "纸质", "规格", "每箱个数", "产品订单数量", "纸箱数量", "单位"]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=9, column=column, value=header)
        cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=teal)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=medium, bottom=medium)
    sheet.row_dimensions[9].height = 28

    first_line_row = 10
    for index, line in enumerate(lines, start=1):
        row = first_line_row + index - 1
        values = [
            index,
            line.packaging_type,
            line.paper_quality,
            f"{line.specification} {line.dimension_unit}".strip(),
            float(line.usage_quantity),
            float(order.product_order_quantity),
            None,
            line.unit,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.font = Font(name="Microsoft YaHei", size=9, color="0F172A")
            cell.alignment = Alignment(
                horizontal="center" if column in {1, 8} else "right" if column in {5, 6, 7} else "left",
                vertical="center",
                wrap_text=column == 4,
            )
            cell.border = Border(left=thin, right=thin, bottom=thin)
            if index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
        sheet.cell(row=row, column=7, value=f"=ROUNDUP(F{row}/E{row},0)")
        sheet.cell(row=row, column=5).number_format = (
            "#,##0" if line.usage_quantity == line.usage_quantity.to_integral_value() else "#,##0.######"
        )
        sheet.cell(row=row, column=6).number_format = (
            "#,##0" if order.product_order_quantity == order.product_order_quantity.to_integral_value() else "#,##0.######"
        )
        sheet.cell(row=row, column=7).number_format = "#,##0"
        sheet.row_dimensions[row].height = 24

    footer_row = first_line_row + len(lines) + 1
    sheet.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=8)
    sheet.cell(row=footer_row, column=1, value=f"备注：{order.note or '无'}")
    sheet.cell(row=footer_row, column=1).font = Font(name="Microsoft YaHei", size=9, color=slate)
    sheet.cell(row=footer_row, column=1).alignment = Alignment(vertical="center", wrap_text=True)
    sheet.cell(row=footer_row, column=1).fill = PatternFill("solid", fgColor=slate_light)
    sheet.row_dimensions[footer_row].height = 30

    generated_row = footer_row + 2
    sheet.merge_cells(start_row=generated_row, start_column=1, end_row=generated_row, end_column=3)
    sheet.cell(row=generated_row, column=1, value=f"制单人：{order.created_by_name}")
    sheet.merge_cells(start_row=generated_row, start_column=4, end_row=generated_row, end_column=8)
    sheet.cell(row=generated_row, column=4, value=f"系统生成时间：{generated_at.strftime('%Y-%m-%d %H:%M')}")
    for cell_ref in (f"A{generated_row}", f"D{generated_row}"):
        sheet[cell_ref].font = Font(name="Microsoft YaHei", size=9, color=slate)

    notice_row = generated_row + 2
    sheet.merge_cells(start_row=notice_row, start_column=1, end_row=notice_row, end_column=8)
    sheet.cell(
        row=notice_row,
        column=1,
        value="本采购单由系统正式订单自动生成，无需供应商回签确认；后续按计划交期进入排期核对、收料和库存流程。",
    )
    sheet.cell(row=notice_row, column=1).font = Font(name="Microsoft YaHei", size=9, bold=True, color=teal_dark)
    sheet.cell(row=notice_row, column=1).fill = PatternFill("solid", fgColor=teal_light)
    sheet.cell(row=notice_row, column=1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[notice_row].height = 32

    widths = [8, 15, 16, 30, 14, 16, 16, 18]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.print_area = f"A1:H{notice_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.35
    sheet.page_margins.bottom = 0.35

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def build_combined_purchase_order_workbook(
    orders: list[tuple[CartonOrder, list[CartonOrderLine]]],
    *,
    generated_at: datetime,
) -> bytes:
    workbook = Workbook()
    workbook.properties.creator = "Royal Regent Nexus"
    workbook.properties.title = f"纸箱累计对账表（{len(orders)} 张订单）"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    sheet = workbook.active
    sheet.title = "纸箱累计对账表"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A8"

    teal = "087F74"
    teal_dark = "075E57"
    teal_light = "E8F6F4"
    slate = "475569"
    slate_light = "F1F5F9"
    border_color = "CBD5E1"
    white = "FFFFFF"
    thin = Side(style="thin", color=border_color)
    medium = Side(style="medium", color=teal)
    factory_id = orders[0][0].factory_id
    supplier_name = orders[0][0].supplier_name_snapshot
    batch_no = f"PO-BATCH-{generated_at.strftime('%Y%m%d%H%M%S')}"
    total_product_quantity = sum(order.product_order_quantity for order, _ in orders)

    sheet.merge_cells("A1:O1")
    sheet["A1"] = f"Royal Regent Nexus · {FACTORY_NAMES.get(factory_id, factory_id)}纸箱累计对账表"
    sheet["A1"].font = Font(name="Microsoft YaHei", size=18, bold=True, color=white)
    sheet["A1"].fill = PatternFill("solid", fgColor=teal_dark)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 34

    sheet.merge_cells("A2:O2")
    sheet["A2"] = "CUMULATIVE PAPER PACKAGING RECONCILIATION"
    sheet["A2"].font = Font(name="Arial", size=9, bold=True, color=teal_dark)
    sheet["A2"].fill = PatternFill("solid", fgColor=teal_light)
    sheet["A2"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[2].height = 21

    information = {
        "A4": "对账批次号",
        "B4": batch_no,
        "E4": "采购厂区",
        "F4": FACTORY_NAMES.get(factory_id, factory_id),
        "H4": "供应商",
        "I4": supplier_name,
        "A5": "系统生成时间",
        "B5": generated_at.strftime("%Y-%m-%d %H:%M"),
        "E5": "订单数量",
        "F5": len(orders),
        "H5": "产品数量合计",
        "I5": float(total_product_quantity),
    }
    for cell_ref, value in information.items():
        sheet[cell_ref] = value
    for merged_range in ("B4:D4", "F4:G4", "I4:O4", "B5:D5", "F5:G5", "I5:O5"):
        sheet.merge_cells(merged_range)
    label_cells = ("A4", "E4", "H4", "A5", "E5", "H5")
    for cell_ref in label_cells:
        cell = sheet[cell_ref]
        cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color=slate)
        cell.fill = PatternFill("solid", fgColor=slate_light)
        cell.alignment = Alignment(horizontal="left", vertical="center")
    for row in sheet.iter_rows(min_row=4, max_row=5, min_col=1, max_col=15):
        for cell in row:
            cell.border = Border(bottom=thin)
            if cell.coordinate not in label_cells:
                cell.font = Font(name="Microsoft YaHei", size=10, color="0F172A")
                cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet["I5"].number_format = (
        "#,##0" if total_product_quantity == total_product_quantity.to_integral_value() else "#,##0.######"
    )

    headers = [
        "序号",
        "采购单号",
        "客户",
        "合同号",
        "货号",
        "产品名称",
        "下单日期",
        "计划交期",
        "产品数量",
        "纸品类型",
        "纸质",
        "规格",
        "每箱个数",
        "纸箱数量",
        "单位",
    ]
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=7, column=column, value=header)
        cell.font = Font(name="Microsoft YaHei", size=9, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=teal)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=medium, bottom=medium)
    sheet.row_dimensions[7].height = 30

    current_row = 8
    line_count = 0
    for order_index, (order, lines) in enumerate(orders, start=1):
        for line_index, line in enumerate(lines, start=1):
            values = [
                order_index,
                order.order_no,
                order.customer_name,
                order.contract_no,
                order.item_no,
                order.product_name or "—",
                order.order_date,
                order.due_date,
                float(order.product_order_quantity),
                line.packaging_type,
                line.paper_quality,
                f"{line.specification} {line.dimension_unit}".strip(),
                float(line.usage_quantity),
                None,
                line.unit,
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=current_row, column=column, value=value)
                cell.font = Font(name="Microsoft YaHei", size=9, color="0F172A")
                cell.alignment = Alignment(
                    horizontal="center" if column in {1, 7, 8, 15} else "right" if column in {9, 13, 14} else "left",
                    vertical="center",
                    wrap_text=column in {3, 4, 5, 6, 12},
                )
                cell.border = Border(left=thin, right=thin, bottom=thin)
                if order_index % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F8FAFC")
            sheet.cell(row=current_row, column=14, value=f"=ROUNDUP(I{current_row}/M{current_row},0)")
            sheet.cell(row=current_row, column=9).number_format = (
                "#,##0" if order.product_order_quantity == order.product_order_quantity.to_integral_value() else "#,##0.######"
            )
            sheet.cell(row=current_row, column=13).number_format = (
                "#,##0" if line.usage_quantity == line.usage_quantity.to_integral_value() else "#,##0.######"
            )
            sheet.cell(row=current_row, column=14).number_format = "#,##0"
            if line_index == 1:
                for column in range(1, 16):
                    sheet.cell(row=current_row, column=column).border = Border(
                        left=thin,
                        right=thin,
                        top=medium,
                        bottom=thin,
                    )
            sheet.row_dimensions[current_row].height = 28
            current_row += 1
            line_count += 1

    summary_row = current_row + 1
    sheet.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=15)
    sheet.cell(
        row=summary_row,
        column=1,
        value=f"合计：{len(orders)} 张订单 · {line_count} 条纸品明细 · 产品数量 {float(total_product_quantity):,.6f}".rstrip("0").rstrip("."),
    )
    sheet.cell(row=summary_row, column=1).font = Font(name="Microsoft YaHei", size=10, bold=True, color=teal_dark)
    sheet.cell(row=summary_row, column=1).fill = PatternFill("solid", fgColor=teal_light)
    sheet.cell(row=summary_row, column=1).alignment = Alignment(horizontal="right", vertical="center")
    sheet.row_dimensions[summary_row].height = 28

    notes = [f"{order.order_no}：{order.note}" for order, _ in orders if order.note]
    notes_row = summary_row + 2
    sheet.merge_cells(start_row=notes_row, start_column=1, end_row=notes_row, end_column=15)
    sheet.cell(row=notes_row, column=1, value="订单备注：" + ("；".join(notes) if notes else "无"))
    sheet.cell(row=notes_row, column=1).font = Font(name="Microsoft YaHei", size=9, color=slate)
    sheet.cell(row=notes_row, column=1).fill = PatternFill("solid", fgColor=slate_light)
    sheet.cell(row=notes_row, column=1).alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[notes_row].height = max(30, 16 * (len(notes) + 1))

    generated_row = notes_row + 2
    creators = "、".join(dict.fromkeys(order.created_by_name for order, _ in orders if order.created_by_name)) or "—"
    sheet.merge_cells(start_row=generated_row, start_column=1, end_row=generated_row, end_column=7)
    sheet.cell(row=generated_row, column=1, value=f"制单人：{creators}")
    sheet.merge_cells(start_row=generated_row, start_column=8, end_row=generated_row, end_column=15)
    sheet.cell(row=generated_row, column=8, value=f"系统生成时间：{generated_at.strftime('%Y-%m-%d %H:%M')}")
    for cell_ref in (f"A{generated_row}", f"H{generated_row}"):
        sheet[cell_ref].font = Font(name="Microsoft YaHei", size=9, color=slate)

    notice_row = generated_row + 2
    sheet.merge_cells(start_row=notice_row, start_column=1, end_row=notice_row, end_column=15)
    sheet.cell(
        row=notice_row,
        column=1,
        value="本文件仅用于核对所选订单的当前累计总量，不代表向供应商首次、追加或减少下单；正式执行请以独立采购单/变更单为准。",
    )
    sheet.cell(row=notice_row, column=1).font = Font(name="Microsoft YaHei", size=9, bold=True, color=teal_dark)
    sheet.cell(row=notice_row, column=1).fill = PatternFill("solid", fgColor=teal_light)
    sheet.cell(row=notice_row, column=1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[notice_row].height = 32

    widths = [8, 22, 18, 22, 18, 24, 13, 13, 14, 14, 16, 28, 14, 14, 10]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.auto_filter.ref = f"A7:O{current_row - 1}"
    sheet.print_area = f"A1:O{notice_row}"
    sheet.print_title_rows = "1:7"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.3
    sheet.page_margins.bottom = 0.3

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
