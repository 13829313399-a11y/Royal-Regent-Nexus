"""Shared, bounded Excel reader for new regular and special orders (no database)."""
import hashlib
import re
from datetime import date, datetime
from io import BytesIO
from zipfile import ZipFile, BadZipFile

from fastapi import HTTPException

MAX_BYTES = 5 * 1024 * 1024
DEMO_COLUMN = "模板行标记"
DEMO_MARKER = "演示行（不导入）"
REQUIREMENT_LABELS = {"required": "必填", "optional": "选填", "conditional": "按条件必填"}
# Ambiguous bare 数量/交期 are deliberately not aliases: the operator maps them.
COMMON = [
    ("group", "分组编号", ["分组", "订单分组"]),
    ("item_no", "货号", ["产品编号", "产品货号", "item", "item no"]),
    ("packaging_type", "纸品类型", ["纸品名称", "包装类型"]),
    ("paper_quality", "纸质", ["材质"]),
    ("specification", "规格", ["尺寸", "纸品规格"]),
    ("dimension_unit", "尺寸单位", []),
    ("quantity", "指定采购数量", ["需求数量", "纸品数量", "采购数量", "纸箱数量"]),
    ("unit", "数量单位", ["单位"]),
    ("note", "备注", ["说明"]),
]
REGULAR = [
    ("customer", "客户", ["客户名称", "客户编码", "customer"]),
    ("contract_no", "合同号", ["合同编号", "contract", "contract no"]),
    ("customer_po", "客户PO", ["客户 PO", "PO", "po no"]),
    ("product_name", "产品名称", ["品名"]),
    ("product_quantity", "产品订单数量", ["产品数量", "订单产品数量"]),
    ("order_date", "下单日期", ["订单日期"]),
    ("customer_due_date", "客户交期", ["客户要求交期"]),
    ("due_date", "计划到料日期", ["计划交期", "计划到货日期"]),
    ("usage_quantity", "每箱个数", ["装箱数", "每箱装量"]),
    ("carton_source_row", "配套纸箱原表行号", ["配套纸箱行号"]),
    ("per_carton_quantity", "每箱配套用量", ["每箱纸卡用量"]),
    ("net_weight_kg", "每箱净重kg", ["每箱净重"]),
    ("gross_weight_kg", "每箱毛重kg", ["每箱毛重"]),
]
SPECIAL = [("kind", "订单类型", ["类型"]), ("required_date", "需用日期", ["需求日期", "需要日期"])]


def fields(mode):
    return COMMON + (REGULAR if mode == "REGULAR" else SPECIAL)


def field_specs(mode):
    required = {"packaging_type", "paper_quality", "specification", "unit"}
    required |= {"customer", "contract_no", "item_no", "order_date", "customer_due_date"} if mode == "REGULAR" else {"group", "kind", "required_date", "quantity"}
    conditional = {"quantity", "product_quantity", "usage_quantity", "carton_source_row", "per_carton_quantity", "customer_po"} if mode == "REGULAR" else set()
    hints = {
        "group": "同一订单全部纸品填写相同分组编号。" if mode == "SPECIAL" else "可用来核对分组；同一编号的客户、合同、货号和 PO 须一致。",
        "item_no": "按文本填写，保留前导零。",
        "packaging_type": "例如外箱、内盒、平卡。", "paper_quality": "例如 A33，按实际纸质填写。",
        "specification": "例如 40*30*25；尺寸单位另列填写。", "dimension_unit": "留空按 cm；使用其他单位时请明确填写。",
        "quantity": "填写每行纸品的明确需求数量。" if mode == "SPECIAL" else "整单直接指定采购数量时，每行必填；按装箱或配套规则计算时可留空。",
        "unit": "例如个、箱、张，须与采购数量对应。", "note": "可填写纸品备注，无备注可留空。",
        "customer": "填写本厂已维护的准确客户名称或编码。", "contract_no": "按文本填写合同号，保留前导零。",
        "customer_po": "客户规则要求 PO 时必填；否则可空白，同一订单各行保持一致。",
        "product_name": "可填写产品名称。",
        "product_quantity": "按装箱或配套规则计算时必填；整单所有纸品直接指定采购数量时可不填。",
        "order_date": "格式 YYYY-MM-DD，同一订单各行一致。",
        "customer_due_date": "格式 YYYY-MM-DD，不早于下单日期；同一订单各行一致。",
        "due_date": "留空按客户交期和当前提前天数计算；手工填写时使用 YYYY-MM-DD。",
        "usage_quantity": "按装箱规则计算的纸箱行必填；整单直接指定采购数量或该行按配套规则计算时可不填。",
        "carton_source_row": "配套纸品按纸箱数量计算时必填，引用同一订单纸箱的实际 Excel 行号，不能引用演示行。",
        "per_carton_quantity": "配套纸品按纸箱数量计算时必填，并同时填写配套纸箱原表行号。",
        "net_weight_kg": "按实际每箱净重填写，单位 kg。", "gross_weight_kg": "按实际每箱毛重填写，单位 kg。",
        "kind": "填写样板单或杂单，也可在预览中统一补齐空白类型。",
        "required_date": "格式 YYYY-MM-DD，也可在预览中统一补齐空白日期。",
    }
    return [{"key": key, "label": label, "requirement": "required" if key in required else "conditional" if key in conditional else "optional", "hint": hints[key]} for key, label, _ in fields(mode)]


def demo_values(mode):
    # Stable template example, not business data. Keep values stable for old downloads.
    values = dict(group="示例-A", item_no="000123", packaging_type="外箱", paper_quality="A33", specification="40*30*25",
                  dimension_unit="cm", quantity=100, unit="个", note="仅演示；请在下一行填写实际订单。")
    if mode == "REGULAR":
        values.update(customer="示例客户", contract_no="DEMO-001", customer_po="", product_name="示例产品", product_quantity="",
                      order_date=date(2026, 11, 1), customer_due_date=date(2026, 11, 20), due_date="", usage_quantity="",
                      carton_source_row="", per_carton_quantity="", net_weight_kg="", gross_weight_kg="")
    else:
        values.update(kind="杂单", required_date=date(2026, 11, 20))
    return values


def normalized(value):
    value = re.sub(r"\s*[（(](?:必填|选填|按条件必填)[）)]\s*$", "", value)
    return re.sub(r"[\s_（）().]+", "", value).casefold()


def cell_text(cell):
    if cell.data_type in {"f", "e"}:
        raise HTTPException(422, f"{cell.coordinate} 含公式或错误值，请粘贴为实际数值后导入")
    value = cell.value
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)) and not isinstance(value, bool) and value == int(value):
        text = str(int(value))
        if re.fullmatch(r"0{2,32}", cell.number_format or ""):
            return text.zfill(len(cell.number_format))
        return text
    text = "" if value is None else str(value).strip()
    if len(text) > 2000:
        raise HTTPException(422, f"{cell.coordinate} 内容超过 2000 字")
    return text


def read(content, mode, sheet_name="", header_row=None, mapping=None):
    from openpyxl import load_workbook
    if mode not in {"REGULAR", "SPECIAL"} or not content or len(content) > MAX_BYTES:
        raise HTTPException(422, "请选择不超过 5 MiB 的 .xlsx 文件")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 500 or sum(x.file_size for x in archive.infolist()) > 20 * 1024 * 1024:
                raise HTTPException(422, "文件展开后过大，请缩小批次")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, ValueError, KeyError, OSError) as exc:
        raise HTTPException(422, "无法读取 Excel，请另存为 .xlsx 后重试") from exc
    try:
        if len(book.sheetnames) > 20:
            raise HTTPException(422, "每次最多 20 个工作表，请拆分文件")
        available = [s.title for s in book if s.sheet_state == "visible"]
        if not available or (sheet_name and sheet_name not in available):
            raise HTTPException(422, "请选择有效的可见工作表")
        aliases = {normalized(label): key for key, label, names in fields(mode) for label in [label, *names]}
        candidates = []
        for name in ([sheet_name] if sheet_name else available):
            sheet = book[name]
            for number, cells in enumerate(sheet.iter_rows(max_row=50, max_col=64), 1):
                score = len({aliases.get(normalized(str(c.value or ""))) for c in cells} - {None})
                candidates.append((score, name, number))
        best = max(candidates, key=lambda x: x[0])
        sheet_name = sheet_name or best[1]
        header_row = header_row or best[2]
        if not 1 <= header_row <= 50:
            raise HTTPException(422, "表头行须在第 1–50 行")
        sheet = book[sheet_name]
        if (sheet.max_column or 0) > 64 or (sheet.max_row or 0) > header_row + 501:
            raise HTTPException(422, "所选工作表最多 64 列、500 行实际数据及 1 行模板演示，请删除多余格式行或拆分文件")
        cells = list(sheet.iter_rows(min_row=header_row, max_row=min(sheet.max_row or header_row, header_row + 501), max_col=sheet.max_column or 1))
        headers = [cell_text(c) for c in cells[0]]
        values = [(header_row + n, [cell_text(c) for c in row]) for n, row in enumerate(cells[1:], 1) if any(c.value is not None for c in row)]
        marker_columns = [i for i, label in enumerate(headers) if label == DEMO_COLUMN]
        ignored_demo_rows = []
        if len(marker_columns) > 1:
            raise HTTPException(422, "模板行标记列重复，请核对表头")
        if marker_columns:
            marker = marker_columns[0]
            example = {key: str(value) for key, value in demo_values(mode).items()}
            retained = []
            for number, row in values:
                if row[marker] != DEMO_MARKER:
                    if row[marker]:
                        raise HTTPException(422, f"第 {number} 行模板标记不明确；实际订单请清空“{DEMO_COLUMN}”")
                    retained.append((number, row))
                    continue
                seen = {}
                unchanged = True
                for i, value in enumerate(row):
                    if i == marker:
                        continue
                    key = aliases.get(normalized(headers[i]))
                    if key:
                        if key in seen or value != example[key]:
                            unchanged = False
                        seen[key] = value
                    elif value:
                        unchanged = False
                if not unchanged or set(seen) != set(example):
                    raise HTTPException(422, f"第 {number} 行仍带演示标记，但内容或列已修改；实际订单请清空“{DEMO_COLUMN}”，再重新上传")
                ignored_demo_rows.append(number)
            values = retained
        if len(ignored_demo_rows) > 1 or len(values) > 500 or (sheet.max_row or header_row) - header_row - len(ignored_demo_rows) > 500:
            raise HTTPException(422, "每次最多 500 行实际数据及 1 行模板演示，请删除多余行或拆分文件")
        inferred = {}
        for key, _, _ in fields(mode):
            matches = [i + 1 for i, value in enumerate(headers) if aliases.get(normalized(value)) == key]
            if len(matches) == 1:
                inferred[key] = matches[0]
        mapping = inferred if mapping is None else mapping
        keys = {k for k, _, _ in fields(mode)}
        if set(mapping) - keys or any(not isinstance(col, int) or not 1 <= col <= len(headers) or col - 1 in marker_columns for col in mapping.values()) or len(set(mapping.values())) != len(mapping):
            raise HTTPException(422, "列匹配无效：同一原表列只能对应一个字段")
        rows = [{"source_row": n, "data": {key: row[mapping[key] - 1] if key in mapping else "" for key in keys}} for n, row in values]
        return {"fingerprint": hashlib.sha256(content).hexdigest(), "sheets": available, "sheet": sheet_name,
                "header_row": header_row, "mapping": mapping, "rows": rows,
                "ignored_demo_rows": ignored_demo_rows, "fields": field_specs(mode),
                "columns": [{"index": i + 1, "label": text or f"第 {i + 1} 列（无表头）", "sample": next((r[i] for _, r in values if r[i]), "")[:100]} for i, text in enumerate(headers) if i not in marker_columns]}
    finally:
        book.close()


def template(mode):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    book = Workbook()
    sheet = book.active
    sheet.title = "正单需求" if mode == "REGULAR" else "新建需求"
    columns = field_specs(mode)
    example = demo_values(mode)
    sheet.append([DEMO_COLUMN] + [f"{field['label']}（{REQUIREMENT_LABELS[field['requirement']]}）" for field in columns])
    sheet.append([DEMO_MARKER] + [example[field['key']] for field in columns])
    for i in range(1, len(columns) + 2):
        sheet.column_dimensions[get_column_letter(i)].width = 22
        cell = sheet.cell(1, i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        demo = sheet.cell(2, i)
        demo.font = Font(color="64748B", italic=True)
        demo.fill = PatternFill("solid", fgColor="F1F5F9")
        demo.alignment = Alignment(vertical="center", wrap_text=True)
        demo.number_format = "yyyy-mm-dd" if isinstance(demo.value, date) else "0" if isinstance(demo.value, (int, float)) else "@"
        sheet.cell(3, i).number_format = "@"
    sheet.row_dimensions[1].height = 42
    sheet.row_dimensions[2].height = 54
    sheet.freeze_panes = "B3"
    instructions = book.create_sheet("填写说明")
    notes = ["必填须补齐；选填可留空；按条件必填请按下方说明选择数量算法或遵循客户规则。",
             "灰色演示行自动跳过，不计入订单、纸品或数量。实际订单从下一行填；复制演示行时须清空“模板行标记”。",
             "支持自有 Excel 和任意列顺序；一行一种纸品，同一订单全部纸品一起导入。不得漏行或用分组编号代替正单身份。",
             "最多 500 行实际数据及 1 行演示；公式请转为实际值，合并单元格不会自动向下填充。"]
    for index, note in enumerate(notes, 1):
        instructions.append([note])
        instructions.merge_cells(start_row=index, start_column=1, end_row=index, end_column=3)
        instructions.row_dimensions[index].height = 32
        instructions.cell(index, 1).alignment = Alignment(wrap_text=True, vertical="center")
    instructions.append(["字段", "填写要求", "填写说明"])
    for field in columns:
        instructions.append([field["label"], REQUIREMENT_LABELS[field["requirement"]], field["hint"]])
    for cell in instructions[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    for row in instructions.iter_rows(min_row=6):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        instructions.row_dimensions[row[0].row].height = 36
    for col, width in {"A": 24, "B": 18, "C": 86}.items():
        instructions.column_dimensions[col].width = width
    instructions.freeze_panes = "A6"
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()
