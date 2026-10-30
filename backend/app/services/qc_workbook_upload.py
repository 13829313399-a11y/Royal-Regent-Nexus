"""Bounded local XLS conversion and template validation; original uploads stay immutable."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile,ZIP_DEFLATED,BadZipFile
import math,os,re,shutil,struct,subprocess
import xlrd,olefile
from lxml import etree as E
from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils import coordinate_to_tuple,get_column_letter
from openpyxl.utils.escape import escape
from app.services import qc_workbook_xml as w

MAX_UPLOAD=32*1024*1024

def package(content):
    if not content.startswith(b'PK'):raise HTTPException(422,'文件不是有效的 Excel 工作簿')
    try:
        with ZipFile(BytesIO(content)) as z:
            infos=z.infolist()
            if len(infos)>10000 or sum(i.file_size for i in infos)>512*1024*1024 or any(i.flag_bits&1 for i in infos):raise ValueError()
            if len({i.filename for i in infos})!=len(infos):raise ValueError()
            for i in infos:
                if i.filename.endswith(('.xml','.rels')):
                    data=z.read(i)
                    if b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():raise ValueError()
            if z.testzip() or 'xl/workbook.xml' not in z.namelist():raise ValueError()
            if 'xl/vbaProject.bin' in z.namelist() or b'macroEnabled' in z.read('[Content_Types].xml'):raise HTTPException(422,'本功能仅支持普通 XLS/XLSX，不支持宏工作簿')
        return w.Package(content)
    except (BadZipFile,ValueError,KeyError,E.XMLSyntaxError,OverflowError) as e:raise HTTPException(422,'工作簿损坏、加密或超出支持范围；当前表格未修改') from e

def pack(parts):
    out=BytesIO()
    with ZipFile(out,'w',ZIP_DEFLATED) as z:
        for name,data in parts.items():z.writestr(name,data)
    return out.getvalue()

def biff_records(content):
    with olefile.OleFileIO(BytesIO(content)) as ole:
        raw=ole.openstream('Workbook' if ole.exists('Workbook') else 'Book').read()
    offset=0;names={};name=None;rows={}
    while offset+4<=len(raw):
        op,size=struct.unpack_from('<HH',raw,offset)
        if not op and not size:break
        data=raw[offset+4:offset+4+size]
        if len(data)!=size:raise ValueError('Truncated BIFF record')
        name=names.get(offset,name)
        if op==0x85:
            start,_,kind,length,wide=struct.unpack_from('<IBBBB',data)
            if kind!=0:raise HTTPException(422,'XLS 含非普通工作表，请先另存为 XLSX 后上传')
            names[start]=data[8:8+length*(2 if wide&1 else 1)].decode('utf-16le' if wide&1 else 'latin1')
        if name:rows.setdefault(name,[]).append((op,data))
        offset+=size+4
    return rows

def preserve_xls(original,converted):
    """Repair evidenced converter differences, then verify cells, formulas and structure."""
    book=xlrd.open_workbook(file_contents=original,formatting_info=True,ragged_rows=True)
    p=package(converted)
    if bool(book.datemode)!=(p.epoch.year==1904):raise HTTPException(422,'转换未保留日期基准')
    records=biff_records(original);receipt={'removedInvalidFilters':[],'fontColorsRestored':0,'commentsRestored':0,'cellsVerified':0}
    if book.sheet_names()!=p.names():raise HTTPException(422,'XLS 转换后的工作表清单不一致，未切换')
    xfs=p.styles.find(w.tag('cellXfs'));fonts=p.styles.find(w.tag('fonts'));fills=p.styles.find(w.tag('fills'));font_sources={};fill_sources={}
    def rgb(index):
        if index in (64,32767):return 'FF000000'
        if index==65:return 'FFFFFFFF'
        color=book.colour_map.get(index)
        if color is None:raise HTTPException(422,'XLS 存在无法确认的颜色，请先另存为 XLSX')
        return 'FF'+''.join(f'{c:02X}' for c in color)
    for src in book.sheets():
        root=p.sheet(src.name);node=next(s for s in p.sheets if s.get('name')==src.name)
        if ['visible','hidden','veryHidden'][src.visibility]!=node.get('state','visible'):raise HTTPException(422,'转换未保留隐藏状态')
        expected={f'{get_column_letter(c+1)}{a+1}:{get_column_letter(d)}{z}' for a,z,c,d in src.merged_cells}
        if expected!={m.get('ref') for m in root.findall(w.tag('mergeCells')+'/'+w.tag('mergeCell'))}:raise HTTPException(422,'转换未保留合并区域')
        values={int(r.get('r')):p.values(r) for r in p.rows(src.name)}
        for ri in range(src.nrows):
            for ci in range(src.row_len(ri)):
                cell=src.cell(ri,ci)
                if cell.ctype in (xlrd.XL_CELL_EMPTY,xlrd.XL_CELL_BLANK):continue
                want=xlrd.error_text_from_code[cell.value] if cell.ctype==xlrd.XL_CELL_ERROR else cell.value
                got=values.get(ri+1,{}).get(ci+1,'')
                equal=want==got or isinstance(want,(int,float)) and isinstance(got,(int,float)) and math.isclose(want,got,rel_tol=1e-12,abs_tol=1e-9)
                if not equal:raise HTTPException(422,f'转换值不一致：{src.name}!{get_column_letter(ci+1)}{ri+1}，原工作簿未替换')
                receipt['cellsVerified']+=1
        for f in root.findall(w.tag('autoFilter')):
            ref=f.get('ref','')
            if '#REF!' in ref or '!' in ref:
                receipt['removedInvalidFilters'].append({'sheet':src.name,'reference':ref});root.remove(f)
        for cell in root.findall('.//'+w.tag('sheetData')+'/'+w.tag('row')+'/'+w.tag('c')):
            ri,ci=coordinate_to_tuple(cell.get('r'));ri-=1;ci-=1
            if ri>=src.nrows or ci>=src.row_len(ri):continue
            xf=book.xf_list[src.cell_xf_index(ri,ci)];style=xfs[int(cell.get('s','0'))]
            font_sources.setdefault(int(style.get('fontId','0')),set()).add(rgb(book.font_list[xf.font_index].colour_index))
            fill_sources.setdefault(int(style.get('fillId','0')),set()).add((xf.background.fill_pattern,xf.background.pattern_colour_index))
        rr=records.get(src.name,[]);formula_refs={get_column_letter(struct.unpack_from('<H',b,2)[0]+1)+str(struct.unpack_from('<H',b)[0]+1) for op,b in rr if op==6}
        converted_refs={c.get('r') for c in root.findall('.//'+w.tag('c')) if c.find(w.tag('f')) is not None}
        if formula_refs!=converted_refs:raise HTTPException(422,'转换后的公式位置不一致：'+src.name)
        objects={};notes=[];current=None
        for i,(op,b) in enumerate(rr):
            if op==0x5d:current=struct.unpack_from('<H',b,6)[0]
            elif op==0x1b6:
                count=struct.unpack_from('<H',b,10)[0];next_op,chunk=rr[i+1]
                if next_op!=0x3c or len(chunk)<1 or len(chunk)!=count*(2 if chunk[0]&1 else 1)+1:raise HTTPException(422,'XLS 批注编码无法完整保留，请先另存为 XLSX')
                objects[current]=chunk[1:].decode('utf-16le' if chunk[0]&1 else 'latin1')
            elif op==0x1c:
                row,col,_,obj=struct.unpack_from('<HHHH',b);notes.append((get_column_letter(col+1)+str(row+1),obj))
        if notes:
            path=p.path(src.name);relpath=str(Path(path).parent/'_rels'/(Path(path).name+'.rels'))
            rels=E.fromstring(p.parts[relpath]);target=next(r.get('Target') for r in rels if r.get('Type')==w.R+'/comments')
            import posixpath
            target=posixpath.normpath(posixpath.dirname(path)+'/'+target);comments=E.fromstring(p.parts[target])
            for ref,obj in notes:
                comment=next(c for c in comments.find(w.tag('commentList')) if c.get('ref')==ref);texts=comment.findall('.//'+w.tag('t'));raw=objects[obj]
                if sum(len(t.text or '') for t in texts)!=len(raw):raise HTTPException(422,'转换未保留完整批注：'+src.name+'!'+ref)
                offset=0
                for t in texts:
                    n=len(t.text or '');t.text=escape(raw[offset:offset+n]).replace('\x00','_x0000_');offset+=n
                receipt['commentsRestored']+=1
            p.parts[target]=w.xml(comments)
    for fid,colors in font_sources.items():
        if len(colors)!=1:raise HTTPException(422,'转换后的字体颜色无法唯一还原，请先另存为 XLSX')
        color=fonts[fid].find(w.tag('color'))
        if color is None:color=E.SubElement(fonts[fid],w.tag('color'))
        color.attrib.clear();color.set('rgb',next(iter(colors)));receipt['fontColorsRestored']+=1
    # Preserve converter-resolved theme/indexed/auto fills. Font automatic color is a separate context.
    receipt['version']=2
    return p.finish(),receipt

def prepare(filename,content,baseline,mime=None):
    if not isinstance(filename,str) or len(filename)>255 or len(content)>MAX_UPLOAD:raise HTTPException(413,'表格文件过大，最多 32 MB')
    suffix=Path(filename).suffix.lower();receipt={'format':suffix}
    allowed={'':None,'application/octet-stream':None,'application/vnd.ms-excel':'.xls','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':'.xlsx'}
    if mime is not None and (not isinstance(mime,str) or mime not in allowed or allowed[mime] not in (None,suffix)):raise HTTPException(422,'文件类型与 Excel 格式不一致')
    if suffix=='.xls':
        if not content.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):raise HTTPException(422,'XLS 文件格式与扩展名不一致')
        try:
            with olefile.OleFileIO(BytesIO(content)) as ole:
                if any(any('vba' in part.lower() for part in entry) for entry in ole.listdir()):raise HTTPException(422,'不支持含宏的 XLS 工作簿')
        except HTTPException:raise
        except Exception as exc:raise HTTPException(422,'XLS 文件损坏、加密或不受支持，当前表格未修改') from exc
        executable=os.getenv('QC_XLS_CONVERTER') or shutil.which('x2t') or '/Applications/ONLYOFFICE.app/Contents/Resources/converter/x2t'
        if not Path(executable).is_file():raise HTTPException(503,'服务器未配置 XLS 转换器，当前工作簿保留；请配置 ONLYOFFICE x2t')
        try:
            with TemporaryDirectory(prefix='qc-workbook-') as tmp:
                src=Path(tmp)/'source.xls';dst=Path(tmp)/'converted.xlsx';src.write_bytes(content)
                subprocess.run([executable,str(src),str(dst)],cwd=tmp,timeout=90,check=True,capture_output=True)
                converted=dst.read_bytes()
            content,details=preserve_xls(content,converted);receipt.update(details)
        except HTTPException:raise
        except Exception as exc:raise HTTPException(422,'XLS 转换或保留校验失败，当前工作簿未修改') from exc
    elif suffix!='.xlsx':raise HTTPException(422,'请选择 .xls 或 .xlsx 表格')
    p=package(content);base=package(baseline)
    clients=p.client_names()
    if not clients:raise HTTPException(422,'未识别到支持的客户表及字段')
    for name in clients:
        a=p.profile(name)
        row=next(r for r in p.rows(name) if int(r.get('r'))==a['templateRows']['header']);av=p.values(row)
        canonical=w.HEAD2 if p.second(name) else w.HEAD1
        allowed=[[w.normalized(v) for v in canonical]]
        if name in base.client_names():
            b=base.profile(name);original=next(r for r in base.rows(name) if int(r.get('r'))==b['templateRows']['header']);bv=base.values(original)
            allowed.append([w.normalized(bv.get(i)) for i in range(1,14)])
        if [w.normalized(av.get(i)) for i in range(1,14)] not in allowed:raise HTTPException(422,'客户表列结构不一致：'+name)
    receipt['clientLayouts']={name:(2 if p.second(name) else 1) for name in clients}
    try:
        checked=load_workbook(BytesIO(content),read_only=True);checked.close()
    except Exception as exc:raise HTTPException(422,'Excel 打开校验失败，当前工作簿未修改') from exc
    return content,receipt
