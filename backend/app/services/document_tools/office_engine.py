"""Native Office extraction and isolated LibreOffice rendering.

No business records or input documents are mutated. Office automation runs in
its own process/profile; imported macros and external link updates are disabled.
"""
from __future__ import annotations

import calendar
import datetime as dt
import json
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.print_settings import PrintArea

from .document_ir import (Block, Cancelled, Cell, DocumentIR, EngineResult, Issue,
                          Page, SourceAnchor, Table, ToolError, result_file)

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def _check(cancelled):
    if cancelled():
        raise Cancelled()


def office_executable() -> str | None:
    configured = os.getenv("DOCUMENT_TOOLS_SOFFICE_PATH") or os.getenv("LIBREOFFICE_PATH")
    if not configured:
        from app.core.config import settings
        configured = getattr(settings, "document_tools_office_command", "soffice")
    configured = shutil.which(configured) or configured
    if configured and Path(configured).is_file():
        return configured
    candidates = [shutil.which("soffice"), shutil.which("libreoffice")]
    if os.name == "nt":
        candidates += [str(Path(base) / "LibreOffice/program/soffice.com")
                       for base in (os.getenv("ProgramFiles", "C:/Program Files"),
                                    os.getenv("ProgramFiles(x86)", "C:/Program Files (x86)"))]
    return next((p for p in candidates if p and Path(p).is_file()), None)


def uno_python() -> str | None:
    from app.core.config import settings
    candidates = [os.getenv("DOCUMENT_TOOLS_UNO_PYTHON"), getattr(settings, "document_tools_uno_python", "/usr/bin/python3"), sys.executable]
    office = office_executable()
    if office:
        candidates.insert(1, str(Path(office).parent / "python.exe"))
    for candidate in dict.fromkeys(p for p in candidates if p):
        if not Path(candidate).is_file():
            continue
        try:
            probe = subprocess.run([candidate, "-c", "import uno"], capture_output=True,
                                   timeout=8, **_spawn_kwargs())
            if probe.returncode == 0:
                return candidate
        except (OSError, subprocess.TimeoutExpired):
            pass
    return None


def _spawn_kwargs():
    return {"creationflags": subprocess.CREATE_NO_WINDOW | subprocess.CREATE_NEW_PROCESS_GROUP} if os.name == "nt" else {"start_new_session": True}


def _terminate(process):
    if process.poll() is not None:
        return
    if os.name == "nt":
        subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                       creationflags=subprocess.CREATE_NO_WINDOW, timeout=15)
    else:
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except ProcessLookupError:
            pass
    process.wait(timeout=15)


def _wait(process, cancelled, timeout=None):
    if timeout is None:
        from app.core.config import settings
        timeout = getattr(settings, "document_tools_office_timeout_seconds", 300)
    deadline = time.monotonic() + timeout
    try:
        while process.poll() is None:
            _check(cancelled)
            if time.monotonic() > deadline:
                raise ToolError("OFFICE_TIMEOUT", "Office 转换超时，请拆分文件后重试")
            time.sleep(0.15)
        _check(cancelled)
        if process.returncode:
            raise ToolError("OFFICE_CONVERSION_FAILED", "Office 引擎转换失败，请检查文件是否损坏或受密码保护")
    finally:
        _terminate(process)


def render_office(path: Path, output: Path, options: dict, work_dir: Path, cancelled) -> dict:
    # Windows' Office profile contains deeply nested extension/cache paths.
    # Keep this ephemeral profile short even when the artifact job path is long.
    with tempfile.TemporaryDirectory(prefix="rr-lo-") as profile_dir:
        return _render_office_in_profile(path, output, options, work_dir, cancelled, Path(profile_dir))


def _render_office_in_profile(path: Path, output: Path, options: dict, work_dir: Path, cancelled, profile: Path) -> dict:
    """Export a copy using a dedicated profile; advanced Calc options use UNO."""
    _check(cancelled)
    executable = office_executable()
    if not executable:
        raise ToolError("OFFICE_UNAVAILABLE", "服务器未安装 LibreOffice，暂时无法渲染 Office 文档", "请配置 LibreOffice 后重试")
    work_dir.mkdir(parents=True, exist_ok=True)
    output.parent.mkdir(parents=True, exist_ok=True)
    is_excel = path.suffix.lower() in (".xlsx", ".xls")
    advanced = is_excel or bool(options.get("password"))
    base = [executable, "--headless", "--nologo", "--nodefault", "--nofirststartwizard",
            f"-env:UserInstallation={profile.resolve().as_uri()}"]
    metadata = {"renderer": "LibreOffice", "isolated_profile": True}
    if advanced:
        interpreter = uno_python()
        if not interpreter:
            raise ToolError("UNO_UNAVAILABLE", "缺少可导入 UNO 的独立 Python 解释器，无法可靠应用工作表和打印设置", "配置 DOCUMENT_TOOLS_UNO_PYTHON 后重试")
        pipe_name = "rr_document_" + uuid4().hex
        connection = f"uno:pipe,name={pipe_name};urp;StarOffice.ComponentContext"
        request = work_dir / ("office-request-" + uuid4().hex + ".json")
        response = work_dir / ("office-response-" + uuid4().hex + ".json")
        request.write_text(json.dumps({"connection": connection, "input": str(path.resolve()),
                                      "output": str(output.resolve()), "options": options,
                                      "response": str(response.resolve())}, ensure_ascii=False), encoding="utf-8")
        try:
            office = subprocess.Popen(base + [f"--accept=pipe,name={pipe_name};urp;StarOffice.ServiceManager"],
                                      stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, **_spawn_kwargs())
            try:
                helper = subprocess.Popen([interpreter, str(Path(__file__).with_name("office_uno_helper.py")), str(request)],
                                          stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, **_spawn_kwargs())
                _wait(helper, cancelled)
                payload = json.loads(response.read_text(encoding="utf-8")) if response.exists() else {}
                if payload.get("error"):
                    raise ToolError(payload.get("code", "OFFICE_CONVERSION_FAILED"), payload["error"])
                metadata.update(payload)
            finally:
                _terminate(office)
        finally:
            request.unlink(missing_ok=True)  # Password never remains in a task artifact.
    else:
        export_dir = work_dir / ("office-export-" + uuid4().hex)
        export_dir.mkdir()
        export_format = {".pdf": "pdf:writer_pdf_Export", ".docx": "docx:Office Open XML Text"}.get(output.suffix, output.suffix.lstrip("."))
        if output.suffix == ".pdf":
            filters = {"UseLosslessCompression": {"type": "boolean", "value": "true"},
                       "ReduceImageResolution": {"type": "boolean", "value": "false"}}
            if options.get("page_selection", "all") != "all":
                filters["PageRange"] = {"type": "string", "value": options["page_selection"]}
            export_format += ":" + json.dumps(filters)
        process = subprocess.Popen(base + ["--convert-to", export_format, "--outdir", str(export_dir), str(path.resolve())],
                                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, **_spawn_kwargs())
        _wait(process, cancelled)
        produced = export_dir / (path.stem + output.suffix)
        if produced.exists():
            shutil.copyfile(produced, output)
    if not output.is_file() or output.stat().st_size == 0:
        raise ToolError("OFFICE_NO_OUTPUT", "Office 没有生成有效输出，请检查源文件和密码")
    if output.suffix == ".pdf":
        from pypdf import PdfReader
        try:
            pages = len(PdfReader(output).pages)
            if not pages:
                raise ValueError("empty")
            metadata["rendered_pages"] = pages
        except Exception as exc:
            raise ToolError("INVALID_OUTPUT", "Office 输出 PDF 无法读取") from exc
    elif output.suffix in (".docx", ".xlsx"):
        if not zipfile.is_zipfile(output):
            raise ToolError("INVALID_OUTPUT", "Office 输出不是有效的 Office 文档")
    return metadata


def _issue(ir, code, message, target=None, source=None, severity="warning"):
    ir.issues.append(Issue(id=f"issue-{len(ir.issues)+1}", code=code, message=message,
                           target_id=target, source=source or SourceAnchor(), severity=severity))


def _xml_text(element):
    parts = []
    for node in element.iter():
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag in (qn("w:br"), qn("w:cr")):
            parts.append("\n")
        elif node.tag == qn("w:tab"):
            parts.append("\t")
    return "".join(parts)


def _word_style(tc):
    style = {}
    for key, query, attr in [("background", "./w:tcPr/w:shd", "fill"),
                             ("alignment", "./w:p/w:pPr/w:jc", "val"),
                             ("color", "./w:p/w:r/w:rPr/w:color", "val")]:
        found = tc.find(query, NS)
        if found is not None:
            style[key] = found.get(qn("w:" + attr))
    style["bold"] = tc.find("./w:p/w:r/w:rPr/w:b", NS) is not None
    borders = tc.find("./w:tcPr/w:tcBorders", NS)
    if borders is not None:
        style["borders"] = {etree.QName(c).localname: {etree.QName(k).localname: v for k, v in c.attrib.items()} for c in borders}
    return style


def _extract_docx(path, options, progress, cancelled):
    document = Document(path)
    ir = DocumentIR(source_type="docx", engine_manifest={"native_reader": "python-docx/OOXML"})
    table_index = 0
    last_title = ""

    def parse_table(element, location, parent_cell=None):
        nonlocal table_index
        table_index += 1
        tid = f"word-table-{table_index}"
        rows = element.findall("w:tr", NS)
        grid = element.findall("w:tblGrid/w:gridCol", NS)
        table = Table(id=tid, title=last_title[:80] or f"表格 {table_index}", row_count=len(rows),
                      column_count=len(grid), source=SourceAnchor(block_id=location, anchor_precision="block"),
                      parent_cell_id=parent_cell, grid_widths_pt=[int(x.get(qn("w:w"), "0")) / 20 for x in grid])
        ir.tables.append(table)
        active = {}
        for r, row in enumerate(rows):
            _check(cancelled)
            before = row.find("w:trPr/w:gridBefore", NS)
            col = int(before.get(qn("w:val"), "0")) if before is not None else 0
            new_active = {}
            if row.find("w:trPr/w:tblHeader", NS) is not None:
                table.header_rows.append(r)
            for tc in row.findall("w:tc", NS):
                span_node = tc.find("w:tcPr/w:gridSpan", NS)
                span = int(span_node.get(qn("w:val"), "1")) if span_node is not None else 1
                merge = tc.find("w:tcPr/w:vMerge", NS)
                continuation = merge is not None and merge.get(qn("w:val"), "continue") == "continue"
                if continuation and col in active and active[col].colspan == span:
                    cell = active[col]
                    cell.rowspan += 1
                    for c in range(col, col+span):
                        new_active[c] = cell
                else:
                    text = "\n".join(_xml_text(p) for p in tc.findall("w:p", NS))
                    cell = Cell(id=f"{tid}-r{r+1}c{col+1}", row=r, column=col, colspan=span,
                                raw_text=text, display_text=text, value=text,
                                source=SourceAnchor(block_id=f"{location}/r{r+1}c{col+1}", anchor_precision="cell"), style=_word_style(tc))
                    table.cells.append(cell)
                    if continuation:
                        _issue(ir, "INVALID_WORD_MERGE", "纵向合并缺少对应起始格，已保留本行内容", cell.id, cell.source)
                    if merge is not None:
                        for c in range(col, col+span):
                            new_active[c] = cell
                    for nested_index, nested in enumerate(tc.findall("w:tbl", NS)):
                        parse_table(nested, cell.source.block_id + f"/nested{nested_index+1}", cell.id)
                col += span
            table.column_count = max(table.column_count, col)
            active = new_active
        return tid

    def parse_container(container, location):
        nonlocal last_title
        for index, element in enumerate(container):
            _check(cancelled)
            source = SourceAnchor(block_id=f"{location}/{index+1}", anchor_precision="block")
            bid = f"word-block-{len(ir.blocks)+1}"
            if element.tag == qn("w:p"):
                text = _xml_text(element)
                pstyle = element.find("w:pPr/w:pStyle", NS)
                pname = pstyle.get(qn("w:val"), "") if pstyle is not None else ""
                kind = "heading" if pname.lower().startswith(("heading", "title")) else "paragraph"
                if text:
                    ir.blocks.append(Block(id=bid, kind=kind, text=text, source=source, style={"word_style": pname, "scope": location}))
                    last_title = text if len(text) < 120 else ""
                if element.find(".//w:drawing", NS) is not None or element.find(".//w:pict", NS) is not None:
                    ir.blocks.append(Block(id=bid+"-image", kind="image", text="图片对象", source=source))
                    _issue(ir, "IMAGE_REQUIRES_REVIEW", "Word 图片未当作原生表格；请在原文核验图片内文字", bid+"-image", source)
            elif element.tag == qn("w:tbl"):
                tid = parse_table(element, source.block_id)
                ir.blocks.append(Block(id=bid, kind="table", table_id=tid, source=source))

    parse_container(document.element.body, "body")
    if options.get("include_headers_footers"):
        seen = set()
        for n, section in enumerate(document.sections):
            for key in ("header", "first_page_header", "even_page_header", "footer", "first_page_footer", "even_page_footer"):
                part = getattr(section, key)
                if part.part.partname not in seen:
                    seen.add(part.part.partname)
                    parse_container(part._element, f"section-{n+1}/{key}")
    with zipfile.ZipFile(path) as archive:
        for name, kind in [("word/footnotes.xml", "footnote"), ("word/endnotes.xml", "endnote"), ("word/comments.xml", "comment")]:
            if name in archive.namelist():
                root = etree.fromstring(archive.read(name), parser=etree.XMLParser(resolve_entities=False, no_network=True))
                for node in root:
                    if node.get(qn("w:type")) in ("separator", "continuationSeparator"):
                        continue
                    value = _xml_text(node)
                    if value:
                        ir.blocks.append(Block(id=f"word-block-{len(ir.blocks)+1}", kind=kind, text=value,
                                               source=SourceAnchor(block_id=f"{kind}/{node.get(qn('w:id'))}", anchor_precision="block")))
        tracked = any(b"<w:ins" in archive.read(name) or b"<w:del" in archive.read(name)
                      for name in archive.namelist() if name == "word/document.xml")
        if tracked:
            _issue(ir, "TRACKED_CHANGES", "文档含修订，结构提取保留可读取文字，请对照原稿确认接受或删除的内容")
    ir.engine_manifest["sections"] = [{"width_pt": s.page_width.pt, "height_pt": s.page_height.pt,
                                      "margin_left_pt": s.left_margin.pt, "margin_right_pt": s.right_margin.pt} for s in document.sections]
    progress("extract", len(ir.blocks), len(ir.blocks))
    return ir


def _display(value, fmt):
    """Format common office values; unsupported formats are surfaced, not guessed."""
    if value is None:
        return "", True
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), True
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        if isinstance(value, dt.time):
            return value.isoformat(), True
        fmt_lower = fmt.lower()
        if "年" in fmt:
            return f"{value.year}年{value.month}月{value.day}日", True
        if "mmm" in fmt_lower:
            return f"{value.day:02d}-{calendar.month_abbr[value.month]}-{value.year}", False
        date_format = "%Y/%m/%d" if "/" in fmt_lower else "%Y-%m-%d"
        return value.strftime(date_format + (" %H:%M:%S" if "h" in fmt_lower else "")), fmt_lower in ("yyyy-mm-dd", "yyyy/mm/dd", "yyyy-mm-dd h:mm:ss", "yyyy-mm-dd hh:mm:ss")
    if not isinstance(value, (int, float)):
        return str(value), True
    if fmt in ("General", "@"):
        return str(value), True
    positive = fmt.split(";")[0]
    clean = re.sub(r"\[[^\]]*\]|_[^ ]|\*.", "", positive).replace('"', "")
    if re.fullmatch(r"0+", clean):
        return str(int(value)).zfill(len(clean)), True
    numeric = re.search(r"[#,0]+(?:\.[0#]+)?", clean)
    if numeric:
        pattern = numeric.group()
        digits = len(pattern.partition(".")[2])
        scaled = value * 100 if "%" in clean else value
        number = format(abs(scaled) if value < 0 and "(" in fmt else scaled, f",.{digits}f" if "," in pattern else f".{digits}f")
        if value < 0 and "(" in fmt:
            number = "(" + number + ")"
        result = clean[:numeric.start()] + number + clean[numeric.end():]
        return result, not bool(re.search(r"[Ee?/]|\[[^\]]*\]", fmt))
    return str(value), False


def _range_for_sheet(workbook, sheet, options):
    requested = options.get("range")
    if requested:
        if requested in workbook.defined_names:
            destinations = list(workbook.defined_names[requested].destinations)
            matches = [coord for title, coord in destinations if title == sheet.title]
            return matches
        if "!" in requested:
            title, requested = requested.rsplit("!", 1)
            if title.strip("'") != sheet.title:
                return []
        return [requested]
    if options.get("print_mode", "original") == "original" and sheet.print_area:
        return [str(r) for r in PrintArea.from_string(str(sheet.print_area)).ranges]
    return [sheet.calculate_dimension()]


def _extract_xlsx(path, options, progress, cancelled):
    workbook = load_workbook(path, data_only=False, keep_links=True)
    cache = load_workbook(path, data_only=True, keep_links=False)
    ir = DocumentIR(source_type="xlsx", engine_manifest={"native_reader": "openpyxl", "formula_results": "cached", "sheets": []})
    selected = options.get("sheets") or [s.title for s in workbook if s.sheet_state == "visible" or options.get("include_hidden")]
    missing = set(selected) - set(workbook.sheetnames)
    if missing:
        raise ToolError("INVALID_SHEETS", "不存在的工作表：" + "、".join(sorted(missing)))
    formula_count = 0
    for sheet_idx, sheet in enumerate(workbook):
        _check(cancelled)
        ir.engine_manifest["sheets"].append({"name": sheet.title, "hidden": sheet.sheet_state != "visible", "rows": sheet.max_row, "columns": sheet.max_column,
                                             "print_area": str(sheet.print_area), "orientation": sheet.page_setup.orientation, "paper_size": sheet.page_setup.paperSize})
        if sheet.title not in selected or (sheet.sheet_state != "visible" and not options.get("include_hidden")):
            continue
        for range_idx, coords in enumerate(_range_for_sheet(workbook, sheet, options)):
            try:
                min_c, min_r, max_c, max_r = range_boundaries(coords)
            except ValueError as exc:
                raise ToolError("INVALID_RANGE", "选区必须是命名区域或 A1:B20 形式的范围") from exc
            if not all((min_c, min_r, max_c, max_r)) or max_c > 16384 or max_r > 1048576:
                raise ToolError("INVALID_RANGE", "请输入有限的单元格范围")
            if (max_c-min_c+1)*(max_r-min_r+1) > 250000:
                raise ToolError("OFFICE_RANGE_TOO_LARGE", "单个选区超过 250000 个单元格，请缩小范围")
            rows = [r for r in range(min_r, max_r+1) if options.get("include_hidden") or not sheet.row_dimensions[r].hidden]
            cols = [c for c in range(min_c, max_c+1) if options.get("include_hidden") or not any(d.hidden and (d.min or 1) <= c <= (d.max or d.min or 1) for d in sheet.column_dimensions.values())]
            if not rows or not cols:
                continue
            rmap, cmap = {r: i for i, r in enumerate(rows)}, {c: i for i, c in enumerate(cols)}
            tid = f"sheet-{sheet_idx+1}-range-{range_idx+1}"
            table = Table(id=tid, title=sheet.title, row_count=len(rows), column_count=len(cols),
                          source=SourceAnchor(sheet=sheet.title, cell=coords, anchor_precision="cell"), header_rows=[0],
                          column_widths=[sheet.column_dimensions[get_column_letter(c)].width or 13 for c in cols],
                          orientation=sheet.page_setup.orientation, paper_size=sheet.page_setup.paperSize)
            if sheet.print_title_rows:
                title_rows = [int(n) for n in re.findall(r"\d+", sheet.print_title_rows)]
                if len(title_rows) == 2:
                    table.header_rows = [rmap[r] for r in range(title_rows[0], title_rows[1]+1) if r in rmap]
            merges = {}
            covered = set()
            for merge in sheet.merged_cells.ranges:
                visible_rows = [r for r in rows if merge.min_row <= r <= merge.max_row]
                visible_cols = [c for c in cols if merge.min_col <= c <= merge.max_col]
                if not visible_rows or not visible_cols:
                    continue
                start = (visible_rows[0], visible_cols[0])
                merges[start] = (len(visible_rows), len(visible_cols), merge.min_row, merge.min_col)
                covered.update((r, c) for r in visible_rows for c in visible_cols if (r, c) != start)
                if start != (merge.min_row, merge.min_col):
                    _issue(ir, "CLIPPED_MERGE", "选区或隐藏项截断合并格，已保留合并格原始值", tid, table.source)
            for r in rows:
                _check(cancelled)
                for c in cols:
                    if (r, c) in covered:
                        continue
                    rs, cs, source_r, source_c = merges.get((r, c), (1, 1, r, c))
                    native = sheet.cell(source_r, source_c)
                    if isinstance(native, MergedCell):
                        continue
                    value = native.value
                    formula = value if native.data_type == "f" else None
                    formula_count += int(formula is not None)
                    source = SourceAnchor(sheet=sheet.title, cell=native.coordinate, anchor_precision="cell")
                    resolution = "resolved"
                    if formula:
                        value = formula if options.get("formula_mode") == "formula" else cache[sheet.title][native.coordinate].value
                        if value is None:
                            resolution = "unresolved"
                    display, exact = _display(value, native.number_format)
                    identifier = f"{tid}-{native.coordinate}"
                    style = {"number_format": native.number_format, "bold": bool(native.font.bold),
                             "font_size": native.font.sz, "alignment": native.alignment.horizontal,
                             "wrap_text": bool(native.alignment.wrap_text), "formula": formula,
                             "cached_value": cache[sheet.title][native.coordinate].value if formula else None,
                             "background": native.fill.fgColor.rgb if native.fill.patternType == "solid" and native.fill.fgColor.type == "rgb" else None,
                             "color": native.font.color.rgb if native.font.color and native.font.color.type == "rgb" else None}
                    if isinstance(value, (dt.date, dt.datetime, dt.time)):
                        serializable = value.isoformat()
                        kind = "date"
                    else:
                        serializable = value
                        kind = "boolean" if isinstance(value, bool) else "number" if isinstance(value, (int, float)) else "text"
                    cell = Cell(id=identifier, row=rmap[r], column=cmap[c], rowspan=rs, colspan=cs,
                                raw_text=str(native.value) if native.value is not None else "", display_text=display,
                                value=serializable, value_kind=kind, resolution=resolution, source=source, style=style)
                    table.cells.append(cell)
                    if formula and resolution == "unresolved":
                        _issue(ir, "FORMULA_CACHE_MISSING", "公式没有缓存结果，需 Calc 重算或选择公式文本；未当作零", identifier, source)
                    if formula and ("[" in formula or "_xlfn." in formula):
                        _issue(ir, "FORMULA_COMPATIBILITY", "公式含外部链接或兼容性函数；外链不会自动更新", identifier, source)
                    if native.data_type == "e":
                        _issue(ir, "FORMULA_ERROR", "源单元格含计算错误：" + display, identifier, source)
                    if not exact:
                        _issue(ir, "CUSTOM_DISPLAY_FORMAT", "复杂显示格式需对照原文，保留原值及格式代码", identifier, source)
                    if kind == "number" and abs(value) >= 1e15:
                        _issue(ir, "SOURCE_NUMERIC_PRECISION", "源数值超过 Excel 15 位精度范围，无法恢复源文件已丢失的数字", identifier, source)
            ir.tables.append(table)
            ir.blocks.append(Block(id=tid+"-block", kind="table", table_id=tid, source=table.source))
        if sheet._charts or sheet._images:
            _issue(ir, "EXCEL_DRAWINGS", "工作表含图表或图片：PDF 使用原工作簿渲染；可编辑 Word 需核对这些对象", source=SourceAnchor(sheet=sheet.title))
        progress("extract", sheet_idx+1, len(workbook.worksheets))
    workbook.close()
    cache.close()
    if not ir.tables:
        raise ToolError("EMPTY_SELECTION", "当前选区没有可见工作表或单元格")
    ir.engine_manifest["formula_count"] = formula_count
    return ir


def _normalize(path, options, work_dir, cancelled):
    path = Path(path)
    if path.suffix.lower() in (".doc", ".xls"):
        import msoffcrypto
        try:
            with path.open("rb") as source:
                encrypted = msoffcrypto.OfficeFile(source)
                if encrypted.is_encrypted() and not options.get("password"):
                    raise ToolError("PASSWORD_REQUIRED", "Office 文件已加密，请输入打开密码")
        except ToolError:
            raise
        except Exception:
            pass  # The actual Office reader gives the format-specific error.
        output = work_dir / ("normalized" + (".docx" if path.suffix.lower() == ".doc" else ".xlsx"))
        render_office(path, output, options, work_dir, cancelled)
        return output
    if not zipfile.is_zipfile(path):
        if path.read_bytes()[:8] == bytes.fromhex("D0CF11E0A1B11AE1"):
            if not options.get("password"):
                raise ToolError("PASSWORD_REQUIRED", "Office 文件已加密，请输入打开密码")
            import msoffcrypto
            output = work_dir / ("unlocked" + path.suffix)
            try:
                with path.open("rb") as source:
                    encrypted = msoffcrypto.OfficeFile(source)
                    encrypted.load_key(password=options["password"], verify_password=True)
                    with output.open("wb") as destination:
                        encrypted.decrypt(destination)
            except Exception as exc:
                output.unlink(missing_ok=True)
                raise ToolError("INVALID_PASSWORD", "密码不正确或加密 Office 文件损坏") from exc
            return output
        raise ToolError("INVALID_OFFICE_FILE", "文件不是可读取的 Office 文档，可能已损坏或扩展名不正确")
    return path


def _extract(path, options, work_dir, progress, cancelled):
    normalized = _normalize(Path(path), options, work_dir, cancelled)
    try:
        ir = _extract_docx(normalized, options, progress, cancelled) if normalized.suffix.lower() == ".docx" else _extract_xlsx(normalized, options, progress, cancelled)
    except (zipfile.BadZipFile, KeyError, ValueError, etree.XMLSyntaxError) as exc:
        raise ToolError("INVALID_OFFICE_FILE", "Office 结构损坏，无法完整读取") from exc
    return normalized, ir


def inspect_office(path, options, work_dir, progress, cancelled):
    work_dir = Path(work_dir)
    work_dir.mkdir(parents=True, exist_ok=True)
    normalized, ir = _extract(path, options, work_dir, progress, cancelled)
    files = []
    preview = work_dir / "source-preview.pdf"
    try:
        progress("rebuild", 0, 1)
        ir.engine_manifest.update(render_office(normalized, preview, options, work_dir, cancelled))
        files.append(result_file(preview, "preview"))
        _load_pages(ir, preview)
        progress("rebuild", 1, 1)
    except Cancelled:
        raise
    except ToolError as exc:
        _issue(ir, exc.code, exc.message)
    return EngineResult(ir, files, {"pages": [p.model_dump() for p in ir.pages],
                                   "sheets": ir.engine_manifest.get("sheets", []), "tables": len(ir.tables)},
                        {"checks": [{"name": "native_structure_readable", "status": "passed", "tables": len(ir.tables)}],
                         "note": "读取成功只证明 Office 原生结构可读，不表示跨软件版式完全一致"})


def _load_pages(ir, path):
    from pypdf import PdfReader
    ir.pages = [Page(page_index=i, display_page_number=i+1, width_pt=float(p.mediabox.width), height_pt=float(p.mediabox.height))
                for i, p in enumerate(PdfReader(path).pages)]


def _validate_pdf_content(ir, output, options):
    from pypdf import PdfReader
    import pdfplumber
    if options.get("page_selection", "all") != "all":
        ir.engine_manifest["content_check"] = {"status": "partial_page_selection", "reason": "原生 Office 未提供稳定分页来源，指定页范围不作全文覆盖比对"}
        return
    visible = ["".join("".join(page.extract_text() or "" for page in PdfReader(output).pages).split())]
    # Font runs can occur in PDF content-stream order rather than geometric
    # reading order. A second native reader avoids flagging intact bold headings.
    with pdfplumber.open(output) as pdf:
        visible.append("".join("".join(page.extract_text() or "" for page in pdf.pages).split()))
    candidates = [(c.id, c.display_text, c.source) for t in ir.tables for c in t.cells if c.display_text and c.resolution != "unresolved"]
    candidates += [(b.id, b.text, b.source) for b in ir.blocks if b.kind in ("paragraph", "heading") and b.text]
    missing = [(identifier, value, source) for identifier, value, source in candidates if not any("".join(value.split()) in text for text in visible)]
    ir.engine_manifest["content_check"] = {"method": "native text whitespace-normalized substring coverage", "checked_units": len(candidates), "missing_units": len(missing)}
    for identifier, value, source in missing[:50]:
        _issue(ir, "RENDERED_CONTENT_MISSING", "PDF 文本层未找到完整原文，可能分页、截断或字体映射问题：" + value[:100], identifier, source)


def _font_report(path, ir):
    """Report fontconfig's actual substitution decisions when available."""
    names = set()
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if (path.suffix == ".docx" and (name == "word/document.xml" or name == "word/styles.xml")):
                root = etree.fromstring(archive.read(name), parser=etree.XMLParser(resolve_entities=False, no_network=True))
                for element in root.findall(".//w:rFonts", NS):
                    names.update(element.get(qn("w:"+key)) for key in ("ascii", "eastAsia", "hAnsi") if element.get(qn("w:"+key)))
            elif path.suffix == ".xlsx" and name == "xl/styles.xml":
                root = etree.fromstring(archive.read(name), parser=etree.XMLParser(resolve_entities=False, no_network=True))
                names.update(element.get("val") for element in root.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}font/{http://schemas.openxmlformats.org/spreadsheetml/2006/main}name") if element.get("val"))
    ir.engine_manifest["requested_fonts"] = sorted(names)
    executable = shutil.which("fc-match")
    if not executable:
        ir.engine_manifest["font_audit"] = {"status": "unavailable", "reason": "此平台未提供 fontconfig 字体匹配清单；未声称与 Microsoft Office 分页一致"}
        return
    substitutions = {}
    for name in sorted(names)[:80]:
        try:
            match = subprocess.run([executable, "--format=%{family}", name], capture_output=True, text=True, timeout=5, **_spawn_kwargs())
            families = match.stdout.strip()
            if match.returncode == 0 and families and name.casefold() not in {x.strip().casefold() for x in families.split(",")}:
                substitutions[name] = families
        except (OSError, subprocess.TimeoutExpired):
            pass
    ir.engine_manifest["font_substitutions"] = substitutions
    for requested, replacement in substitutions.items():
        _issue(ir, "FONT_SUBSTITUTION", f"缺少字体 {requested}；fontconfig 匹配替代为 {replacement}，分页可能变化")


def _corrected_copy(path, ir, output):
    """Patch only addressed XML text/cells, preserving charts and unrelated ZIP parts."""
    changes = {}
    corrected_cells = [c for t in ir.tables for c in t.cells if c.source.method == "manual" or c.resolution == "manually_confirmed"]
    corrected_blocks = [b for b in ir.blocks if b.source.method == "manual"]
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    with zipfile.ZipFile(path) as archive:
        def xml(name):
            if name not in changes:
                changes[name] = etree.fromstring(archive.read(name), parser=parser)
            return changes[name]

        if path.suffix.lower() == ".xlsx":
            namespace = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            main_ns = namespace["m"]
            workbook = xml("xl/workbook.xml")
            relations = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"), parser=parser)
            targets = {r.get("Id"): r.get("Target") for r in relations}
            sheets = {}
            for node in workbook.findall("m:sheets/m:sheet", namespace):
                rid = node.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                target = targets[rid]
                sheets[node.get("name")] = target.lstrip("/") if target.startswith("/") else "xl/" + target
            for cell in corrected_cells:
                if cell.source.sheet not in sheets or not cell.source.cell:
                    raise ToolError("REVISION_ANCHOR_MISSING", "修订单元格缺少准确的工作表来源")
                root = xml(sheets[cell.source.sheet])
                found = root.xpath("m:sheetData/m:row/m:c[@r=$coord]", namespaces=namespace, coord=cell.source.cell)
                if not found:
                    raise ToolError("REVISION_ANCHOR_MISSING", "原文件已找不到被修订单元格：" + cell.source.cell)
                native = found[0]
                for child in list(native):
                    if etree.QName(child).localname in ("v", "f", "is"):
                        native.remove(child)
                text = cell.display_text
                if cell.value_kind in ("number", "integer", "decimal") and re.fullmatch(r"-?\d+(?:\.\d+)?", text) and len(text.replace("-", "").replace(".", "")) <= 15:
                    native.set("t", "n")
                    etree.SubElement(native, "{"+main_ns+"}v").text = text
                else:
                    native.set("t", "inlineStr")
                    inline = etree.SubElement(native, "{"+main_ns+"}is")
                    value = etree.SubElement(inline, "{"+main_ns+"}t")
                    value.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    value.text = text
            calc = workbook.find("m:calcPr", namespace)
            if calc is not None:
                calc.set("fullCalcOnLoad", "1")
                calc.set("forceFullCalc", "1")
        elif path.suffix.lower() == ".docx":
            document = Document(path)
            for item in corrected_cells + corrected_blocks:
                anchor = item.source.block_id or ""
                pieces = anchor.split("/")
                if pieces[0] == "body":
                    container = xml("word/document.xml").find("w:body", NS)
                    pieces = pieces[1:]
                elif pieces[0].startswith("section-") and len(pieces) >= 3:
                    section = document.sections[int(pieces[0].split("-")[1])-1]
                    part = getattr(section, pieces[1]).part
                    container = xml(str(part.partname).lstrip("/"))
                    pieces = pieces[2:]
                elif pieces[0] in ("footnote", "endnote", "comment") and len(pieces) == 2:
                    root = xml("word/" + {"footnote": "footnotes.xml", "endnote": "endnotes.xml", "comment": "comments.xml"}[pieces[0]])
                    found = [node for node in root if node.get(qn("w:id")) == pieces[1]]
                    if not found:
                        raise ToolError("REVISION_ANCHOR_MISSING", "原文件已找不到被修订的脚注或批注")
                    container = found[0]
                    pieces = []
                else:
                    raise ToolError("REVISION_ANCHOR_MISSING", "此 Word 区块没有可直接修订的原生来源")
                for segment in pieces:
                    if segment.isdigit():
                        container = container[int(segment)-1]
                    elif re.fullmatch(r"r\d+c\d+", segment):
                        row_num, column_num = map(int, re.findall(r"\d+", segment))
                        row = container.findall("w:tr", NS)[row_num-1]
                        logical_column = 1
                        found = None
                        for native in row.findall("w:tc", NS):
                            span = native.find("w:tcPr/w:gridSpan", NS)
                            width = int(span.get(qn("w:val"), "1")) if span is not None else 1
                            if logical_column <= column_num < logical_column+width:
                                found = native
                                break
                            logical_column += width
                        if found is None:
                            raise ToolError("REVISION_ANCHOR_MISSING", "无法定位 Word 原生表格单元格")
                        container = found
                    elif segment.startswith("nested"):
                        container = container.findall("w:tbl", NS)[int(segment[6:])-1]
                value = item.display_text if isinstance(item, Cell) else item.text
                texts = container.findall(".//w:t", NS) if container.tag == qn("w:p") else container.findall("w:p//w:t", NS)
                if texts:
                    texts[0].text = value
                    texts[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    for node in texts[1:]:
                        node.text = ""
                else:
                    paragraph = container if container.tag == qn("w:p") else container.find("w:p", NS)
                    if paragraph is None:
                        paragraph = etree.SubElement(container, qn("w:p"))
                    run = etree.SubElement(paragraph, qn("w:r"))
                    etree.SubElement(run, qn("w:t")).text = value
        else:
            raise ToolError("REVISION_FORMAT_UNSUPPORTED", "旧格式需先规范化后再修订")
        with zipfile.ZipFile(output, "w") as destination:
            for info in archive.infolist():
                data = etree.tostring(changes[info.filename], xml_declaration=True, encoding="UTF-8", standalone=True) if info.filename in changes else archive.read(info.filename)
                destination.writestr(info, data)
    return output


def _apply_calc_metadata(ir, metadata, recalculated=None, formula_mode="display"):
    displays = metadata.pop("display_cells", {})
    cell_map = {(c.source.sheet, c.source.cell): c for t in (recalculated.tables if recalculated else []) for c in t.cells}
    resolved = set()
    for table in ir.tables:
        for cell in table.cells:
            if cell.resolution == "manually_confirmed" or cell.source.method == "manual":
                continue
            if formula_mode == "formula" and cell.style.get("formula"):
                continue
            if not cell.source.cell:
                continue
            min_c, min_r, _, _ = range_boundaries(cell.source.cell)
            display = displays.get(cell.source.sheet, {}).get(f"{min_r}:{min_c}")
            if not display:
                continue
            if display.get("error"):
                _issue(ir, "CALC_FORMULA_ERROR", "Calc 计算返回错误：" + display["display"], cell.id, cell.source)
                cell.resolution = "unresolved"
                continue
            fresh = cell_map.get((cell.source.sheet, cell.source.cell))
            if fresh and fresh.value is not None and cell.style.get("formula"):
                if cell.value is not None and cell.value != fresh.value:
                    _issue(ir, "FORMULA_RECALC_CHANGED", f"Calc 重算与原缓存不同：{cell.display_text} → {display['display']}", cell.id, cell.source)
                cell.value, cell.value_kind = fresh.value, fresh.value_kind
                cell.style["formula_result_source"] = "Calc recalculated working copy"
            elif cell.style.get("formula") and "value" in display:
                cell.value, cell.value_kind = display["value"], display["value_kind"]
                cell.style["formula_result_source"] = "Calc working-copy display"
            cell.display_text = display["display"]
            cell.resolution = "resolved"
            resolved.add(cell.id)
    ir.issues = [issue for issue in ir.issues if not (issue.target_id in resolved and issue.code in ("FORMULA_CACHE_MISSING", "CUSTOM_DISPLAY_FORMAT"))]
    graphics = metadata.pop("graphics", [])
    for index, graphic in enumerate(graphics):
        ir.blocks.append(Block(id=f"calc-image-{index+1}", kind="image", source=SourceAnchor(sheet=graphic["sheet"], anchor_precision="block"),
                               style={"image_path": graphic["path"], "width_pt": graphic["width_pt"], "height_pt": graphic["height_pt"]}))
    if graphics and not metadata.get("graphics_failed"):
        ir.issues = [issue for issue in ir.issues if issue.code != "EXCEL_DRAWINGS"]
    metadata["exported_graphics_count"] = len(graphics)
    ir.engine_manifest.update(metadata)


def convert_office(path, operation, options, work_dir, progress, cancelled, ir=None):
    work_dir = Path(work_dir)
    work_dir.mkdir(parents=True, exist_ok=True)
    _check(cancelled)
    supplied_ir = ir is not None
    if ir is None:
        normalized, ir = _extract(path, options, work_dir, progress, cancelled)
    else:
        normalized = _normalize(Path(path), options, work_dir, cancelled)
    if operation in ("word_to_pdf", "excel_to_pdf"):
        if supplied_ir and (any(c.source.method == "manual" or c.resolution == "manually_confirmed" for t in ir.tables for c in t.cells) or any(b.source.method == "manual" for b in ir.blocks)):
            normalized = _corrected_copy(normalized, ir, work_dir / ("corrected" + normalized.suffix))
            options = {**options, "recalculate": True}
        output = work_dir / "output.pdf"
        _font_report(normalized, ir)
        progress("rebuild", 0, 1)
        metadata = render_office(normalized, output, options, work_dir, cancelled)
        if operation == "excel_to_pdf":
            _apply_calc_metadata(ir, metadata, formula_mode=options.get("formula_mode", "display"))
        else:
            ir.engine_manifest.update(metadata)
        _load_pages(ir, output)
        _validate_pdf_content(ir, output, options)
        progress("rebuild", 1, 1)
    elif operation == "word_to_excel":
        output = write_xlsx(ir, work_dir / "output.xlsx", options)
    elif operation == "excel_to_word":
        if ir.engine_manifest.get("formula_count"):
            _issue(ir, "FORMULA_STATIC_OUTPUT", "Word 中公式以结果或公式文本呈现，不继续执行 Excel 计算", severity="info")
        if not supplied_ir and office_executable():
            try:
                recalculated = work_dir / "recalculated.xlsx"
                needs_recalc = any(i.code == "FORMULA_CACHE_MISSING" for i in ir.issues) and options.get("formula_mode", "display") == "display"
                progress("rebuild", 0, 1)
                meta = render_office(normalized, recalculated, {**options, "recalculate": needs_recalc}, work_dir, cancelled)
                fresh = _extract_xlsx(recalculated, options, progress, cancelled)
                _apply_calc_metadata(ir, meta, fresh, options.get("formula_mode", "display"))
                progress("rebuild", 1, 1)
            except Cancelled:
                raise
            except ToolError as exc:
                _issue(ir, exc.code, "Calc 显示值或图表读取未完成：" + exc.message)
        output = write_docx(ir, work_dir / "output.docx", options)
    else:
        raise ToolError("UNSUPPORTED_OPERATION", "该 Office 转换方向不受支持")
    _check(cancelled)
    progress("validate", 1, 1)
    files = [result_file(output)]
    if output.suffix == ".docx":
        try:
            preview = work_dir / "result-preview.pdf"
            ir.engine_manifest.update(render_office(output, preview, {}, work_dir, cancelled))
            files.append(result_file(preview, "preview"))
        except Cancelled:
            raise
        except ToolError as exc:
            _issue(ir, exc.code, "输出校样未生成：" + exc.message)
    checks = [{"name": "output_container_readable", "status": "passed", "format": output.suffix.lstrip(".")}]
    content_check = ir.engine_manifest.get("content_check")
    if content_check and "missing_units" in content_check:
        checks.append({"name": "native_text_coverage", "status": "passed" if not content_check["missing_units"] else "needs_review", **content_check})
    checks.append({"name": "pixel_layout_equivalence", "status": "not_evaluated", "reason": "未声称与原 Microsoft Office 分页逐像素等价"})
    return EngineResult(ir, files, {"tables": len(ir.tables), "pages": len(ir.pages),
                                   "cells": sum(len(t.cells) for t in ir.tables)},
                        {"checks": checks, "note": "检查状态只对应列出的实际检查，不代表所有格式完全无损"})


def _safe_sheet_name(title, names):
    stem = re.sub(r"[\\/*?:\[\]]", "_", title).strip(" '")[:31] or "表格"
    name, n = stem, 1
    while name.casefold() in {x.casefold() for x in names}:
        n += 1
        tail = f" ({n})"
        name = stem[:31-len(tail)] + tail
    return name


def _color(value):
    return str(value)[-6:] if value and re.fullmatch(r"[0-9a-fA-F]{6,8}", str(value)) else None


def _word_output_tables(ir, options):
    if not options.get("merge_continuation_tables") or ir.source_type != "docx":
        return ir.tables
    output = []
    block_order = {b.table_id: i for i, b in enumerate(ir.blocks) if b.table_id}
    previous_id = None
    merged_count = 0
    for original in ir.tables:
        previous = output[-1] if output else None
        headers = sorted(original.header_rows)
        header_count = len(headers)
        signature = lambda t: [(c.row, c.column, c.rowspan, c.colspan, c.display_text) for c in t.cells if c.row in t.header_rows]
        same_context = (previous is not None and original.title == previous.title and
                        original.column_count == previous.column_count and
                        getattr(original, "grid_widths_pt", None) == getattr(previous, "grid_widths_pt", None))
        consecutive = previous_id in block_order and original.id in block_order and block_order[original.id] == block_order[previous_id]+1
        if (same_context and consecutive and header_count and headers == list(range(header_count)) and
                signature(original) == signature(previous) and not getattr(original, "parent_cell_id", None) and
                not any(c.row < header_count < c.row+c.rowspan for c in original.cells)):
            offset = previous.row_count-header_count
            for cell in original.cells:
                if cell.row >= header_count:
                    clone = cell.model_copy(deep=True)
                    clone.row += offset
                    previous.cells.append(clone)
            previous.row_count += original.row_count-header_count
            previous.source_pages = sorted(set(previous.source_pages+original.source_pages))
            merged_count += 1
        else:
            output.append(original.model_copy(deep=True))
        previous_id = original.id
    if merged_count:
        ir.engine_manifest["continuation_tables_merged"] = merged_count
    if len(output) > 1:
        _issue(ir, "CONTINUATION_REVIEW", "无法同时确认表头、列宽、标题上下文和相邻关系的表格仍独立输出")
    return output


def write_xlsx(ir: DocumentIR, path: Path, options: dict) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    thin = Side(style="thin", color="D9D9D9")
    for table in _word_output_tables(ir, options):
        sheet = workbook.create_sheet(_safe_sheet_name(table.title or table.id, workbook.sheetnames))
        for cell in table.cells:
            target = sheet.cell(cell.row+1, cell.column+1)
            value = cell.value
            if cell.resolution == "unresolved":
                value = cell.display_text or "[待核验]"
            if cell.value_kind in ("number", "integer", "decimal", "currency", "percentage") and value is not None:
                try:
                    # Excel cannot represent more than 15 significant numeric digits.
                    numeric_text = str(value).replace("-", "").replace(".", "").lstrip("0")
                    if len(numeric_text) <= 15:
                        value = float(value) if "." in str(value) else int(value)
                    else:
                        value = cell.display_text or str(value)
                except (ValueError, TypeError):
                    value = cell.display_text or str(value)
            elif cell.value_kind == "boolean":
                value = bool(value)
            else:
                value = cell.display_text if cell.display_text else ("" if value is None else str(value))
            target.value = value
            if isinstance(value, str):
                target.data_type = "s"  # Untrusted text beginning '=' never becomes a formula.
            target.number_format = cell.style.get("number_format", "General") if not isinstance(value, str) else "@"
            target.font = Font(name="Calibri", size=11, bold=bool(cell.style.get("bold")), color=_color(cell.style.get("color")) or "172B4D")
            target.alignment = Alignment(horizontal={"start": "left", "end": "right", "both": "justify"}.get(cell.style.get("alignment"), cell.style.get("alignment")) or "left", vertical="top", wrap_text=True)
            target.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            background = _color(cell.style.get("background"))
            if background:
                target.fill = PatternFill("solid", fgColor=background)
            if options.get("preserve_merges", True) and (cell.rowspan > 1 or cell.colspan > 1):
                sheet.merge_cells(start_row=cell.row+1, end_row=cell.row+cell.rowspan,
                                  start_column=cell.column+1, end_column=cell.column+cell.colspan)
            ir.mappings.append({"target_id": cell.id, "output_sheet": sheet.title,
                                "output_cell": target.coordinate, "source": cell.source.model_dump()})
        for column in range(1, table.column_count+1):
            widths = getattr(table, "column_widths", [])
            sheet.column_dimensions[get_column_letter(column)].width = max(10, min(45, widths[column-1] if column <= len(widths) else 20))
        sheet.freeze_panes = "A2" if table.header_rows else None
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    notes = ir.blocks if options.get("word_mode") == "structure" else [b for b in ir.blocks if b.kind != "table"]
    if options.get("include_notes_sheet", True) and notes:
        sheet = workbook.create_sheet(_safe_sheet_name("全文结构" if options.get("word_mode") == "structure" else "原文说明", workbook.sheetnames))
        sheet.append(["顺序", "区块类型", "原文内容", "来源"])
        for index, block in enumerate(notes):
            sheet.append([index+1, block.kind, block.text or block.table_id or "", json.dumps(block.source.model_dump(exclude_none=True), ensure_ascii=False)])
            for c in sheet[sheet.max_row]:
                if isinstance(c.value, str):
                    c.data_type = "s"
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for column, width in {"A": 10, "B": 18, "C": 75, "D": 45}.items():
            sheet.column_dimensions[column].width = width
    elif notes:
        _issue(ir, "NOTES_EXCLUDED", "已按选项排除非表格说明文字；原文及结构记录仍可核验")
    if not workbook.worksheets:
        sheet = workbook.create_sheet("原文说明")
        sheet.append(["未检测到原生表格，请在原文对照中检查图片或文本内容。"])
        _issue(ir, "NO_NATIVE_TABLES", "未检测到原生表格；没有将正文挤入伪造表格")
    workbook.save(path)
    check = load_workbook(path, read_only=True)
    check.close()
    return path


def write_docx(ir: DocumentIR, path: Path, options: dict) -> Path:
    if options.get("layout_mode") == "layout" and ir.source_type == "pdf" and ir.pages:
        return _write_positioned_docx(ir, Path(path), options)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    table_map = {t.id: t for t in ir.tables}
    blocks = ir.blocks or [Block(id=t.id+"-block", kind="table", table_id=t.id, source=t.source) for t in ir.tables]
    written_tables = set()
    current_page = None
    first = True

    def page_setup(section, table=None, page=None):
        paper = options.get("paper", "original")
        if paper == "original" and table:
            paper = {"8": "A3", "9": "A4"}.get(str(getattr(table, "paper_size", "")), "original")
        if paper == "A3":
            section.page_width, section.page_height = Inches(11.69), Inches(16.54)
        elif paper == "A4":
            section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        elif page:
            section.page_width, section.page_height = Pt(page.width_pt), Pt(page.height_pt)
        section.left_margin = section.right_margin = Inches(.6)
        section.top_margin = section.bottom_margin = Inches(.6)
        landscape = options.get("orientation") == "landscape" or (options.get("orientation", "auto") == "auto" and table is not None and (table.column_count > 6 or getattr(table, "orientation", None) == "landscape"))
        if landscape and section.page_width < section.page_height:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width

    page_setup(document.sections[0])

    def add_table(native):
        if native.id in written_tables:
            return
        written_tables.add(native.id)
        if native.row_count < 1 or native.column_count < 1:
            return
        if native.row_count * native.column_count > 50000:
            raise ToolError("WORD_TABLE_TOO_LARGE", "Word 表格超过 50000 个单元格，请缩小选区或按工作表转换")
        if native.column_count > 12:
            _issue(ir, "WIDE_TABLE", "表格超过 12 列，已使用横向纸张且保持 10 磅字体；请考虑缩小选区或按列分组", native.id, native.source)
        title = document.add_paragraph(native.title or native.id, style="Heading 2")
        title.paragraph_format.keep_with_next = True
        table = document.add_table(rows=native.row_count, cols=native.column_count)
        table.style = "Table Grid"
        table.autofit = False
        widths = getattr(native, "column_widths", None) or [1]*native.column_count
        available = document.sections[-1].page_width-document.sections[-1].left_margin-document.sections[-1].right_margin
        for index, column in enumerate(table.columns):
            column.width = int(available*widths[index]/sum(widths))
            for cell in column.cells:
                cell.width = column.width
        for row_num in native.header_rows:
            if 0 <= row_num < len(table.rows):
                header = OxmlElement("w:tblHeader")
                table.rows[row_num]._tr.get_or_add_trPr().append(header)
        for cell in native.cells:
            if cell.row >= native.row_count or cell.column >= native.column_count:
                continue
            target = table.cell(cell.row, cell.column)
            if cell.rowspan > 1 or cell.colspan > 1:
                target = target.merge(table.cell(min(native.row_count-1, cell.row+cell.rowspan-1), min(native.column_count-1, cell.column+cell.colspan-1)))
            target.text = cell.display_text if cell.resolution != "unresolved" else (cell.display_text or "[待核验：无公式缓存]")
            for paragraph in target.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                                       "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}.get(cell.style.get("alignment"))
                for run in paragraph.runs:
                    run.font.size, run.bold = Pt(10), bool(cell.style.get("bold"))
                    color = _color(cell.style.get("color"))
                    if color:
                        run.font.color.rgb = RGBColor.from_string(color)
            background = _color(cell.style.get("background"))
            if background:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), background)
                target._tc.get_or_add_tcPr().append(shade)
            ir.mappings.append({"target_id": cell.id, "output_table": native.id, "row": cell.row,
                                "column": cell.column, "source": cell.source.model_dump()})

    for block in blocks:
        page_idx = block.source.page_index
        native = table_map.get(block.table_id) if block.kind == "table" else None
        if (ir.source_type == "xlsx" and native and not first) or (options.get("layout_mode") == "layout" and page_idx is not None and page_idx != current_page and not first):
            document.add_section(WD_SECTION.NEW_PAGE)
        page = next((p for p in ir.pages if p.page_index == page_idx), None)
        if first or native or page_idx != current_page:
            page_setup(document.sections[-1], native, page if options.get("layout_mode") == "layout" else None)
        current_page, first = page_idx, False
        if native:
            add_table(native)
        elif block.kind == "image" and block.style.get("image_path"):
            image_path = Path(block.style["image_path"])
            if image_path.is_file():
                available_width_pt = (document.sections[-1].page_width-document.sections[-1].left_margin-document.sections[-1].right_margin) / 12700
                document.add_picture(str(image_path), width=Pt(min(float(block.style.get("width_pt", 400)), available_width_pt)))
        elif block.text:
            style = "Heading 1" if block.kind == "heading" else "Normal"
            paragraph = document.add_paragraph(block.text, style=style)
            if block.style.get("font_size"):
                for run in paragraph.runs:
                    run.font.size = Pt(max(8, min(36, float(block.style["font_size"]))))
            ir.mappings.append({"target_id": block.id, "output_block": len(document.paragraphs)-1, "source": block.source.model_dump()})
    for native in ir.tables:
        if native.id not in written_tables:
            add_table(native)
    document.save(path)
    Document(path)
    return path


def _write_positioned_docx(ir, path, options):
    """Use editable Word text boxes at PDF point coordinates, never page screenshots.

    VML text boxes are part of OOXML's transitional format and support paragraphs
    and native tables. Rendering remains subject to font and line-metric changes.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.space_after = Pt(0)
    tables = {t.id: t for t in ir.tables}
    vml = "urn:schemas-microsoft-com:vml"
    for page_number, page in enumerate(ir.pages):
        section = document.sections[0] if page_number == 0 else document.add_section(WD_SECTION.NEW_PAGE)
        section.page_width, section.page_height = Pt(page.width_pt), Pt(page.height_pt)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Pt(0)
        anchor = document.add_paragraph()
        anchor.paragraph_format.space_before = anchor.paragraph_format.space_after = Pt(0)
        anchor.paragraph_format.line_spacing = Pt(1)
        blocks = [b for b in ir.blocks if b.source.page_index == page.page_index]
        for index, block in enumerate(blocks):
            box = block.source.bbox_pt
            native = tables.get(block.table_id)
            if not box and native:
                box = native.source.bbox_pt
            if not box:
                _issue(ir, "LAYOUT_ANCHOR_MISSING", "区块缺少原页坐标，放入页面末尾可编辑段落", block.id, block.source)
                document.add_paragraph(block.text)
                continue
            x0, y0, x1, y1 = map(float, box)
            if x1 <= x0 or y1 <= y0:
                _issue(ir, "LAYOUT_INVALID_BOX", "区块坐标无效，需核验来源区域", block.id, block.source)
                continue
            picture = OxmlElement("w:pict")
            shape = etree.SubElement(picture, "{"+vml+"}rect", nsmap={"v": vml})
            shape.set("id", f"RRTextBox{page_number}_{index}")
            shape.set("filled", "f")
            shape.set("stroked", "f")
            # Permit small font-metric growth while retaining the source origin.
            shape.set("style", f"position:absolute;margin-left:{x0:.3f}pt;margin-top:{y0:.3f}pt;width:{x1-x0:.3f}pt;height:{y1-y0+3:.3f}pt;z-index:{index+1};mso-position-horizontal-relative:page;mso-position-vertical-relative:page;v-text-anchor:top")
            textbox = etree.SubElement(shape, "{"+vml+"}textbox")
            textbox.set("inset", "0,0,0,0")
            textbox.set("style", "mso-fit-shape-to-text:t")
            content = etree.SubElement(textbox, qn("w:txbxContent"))
            if native:
                if native.row_count * native.column_count > 50000:
                    raise ToolError("WORD_TABLE_TOO_LARGE", "定位表格超过 50000 个单元格，请缩小选区")
                table = document.add_table(rows=max(1, native.row_count), cols=max(1, native.column_count))
                table.style = "Table Grid"
                table.autofit = False
                for column in table.columns:
                    column.width = Pt((x1-x0)/max(1, native.column_count))
                for cell in native.cells:
                    target = table.cell(cell.row, cell.column)
                    if cell.rowspan > 1 or cell.colspan > 1:
                        target = target.merge(table.cell(cell.row+cell.rowspan-1, cell.column+cell.colspan-1))
                    target.text = cell.display_text or ("[待核验]" if cell.resolution == "unresolved" else "")
                    for paragraph in target.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(0)
                        for run in paragraph.runs:
                            run.font.size = Pt(float(cell.style.get("font_size") or 10))
                    ir.mappings.append({"target_id": cell.id, "output_table": native.id, "row": cell.row,
                                        "column": cell.column, "source": cell.source.model_dump()})
                table._tbl.getparent().remove(table._tbl)
                content.append(table._tbl)
                content.append(OxmlElement("w:p"))
            elif block.kind == "image" and block.style.get("image_path"):
                image_path = Path(block.style["image_path"])
                if image_path.is_file():
                    document.add_picture(str(image_path), width=Pt(x1-x0), height=Pt(y1-y0))
                    paragraph = document.paragraphs[-1]._p
                    paragraph.getparent().remove(paragraph)
                    content.append(paragraph)
            else:
                paragraph = document.add_paragraph(block.text)
                paragraph.paragraph_format.space_after = Pt(0)
                size = float(block.style.get("font_size") or min(12, max(8, (y1-y0)*.85)))
                for run in paragraph.runs:
                    run.font.size = Pt(size)
                    run.bold = bool(block.style.get("bold"))
                paragraph._p.getparent().remove(paragraph._p)
                content.append(paragraph._p)
            anchor.add_run()._r.append(picture)
            ir.mappings.append({"target_id": block.id, "output_page": page.page_index,
                                "output_textbox": shape.get("id"), "bbox_pt": box, "source": block.source.model_dump()})
    ir.engine_manifest["docx_layout"] = {"mode": "editable_positioned_textboxes", "coordinate_unit": "point", "origin": "page-top-left"}
    document.save(path)
    Document(path)
    return path
