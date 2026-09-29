"""Check generated containers with independent readers; no accuracy claims."""
from pathlib import Path
from zipfile import ZipFile

from app.services.document_tools.document_ir import ToolError


def validate_outputs(files):
    checks = []
    for item in files:
        if item["role"] not in {"result", "preview", "package", "result_page", "source_page"}:
            continue
        path = Path(item["path"])
        try:
            if path.suffix.lower() == ".pdf":
                from pypdf import PdfReader
                document = PdfReader(path)
                if not document.pages or any(float(p.mediabox.width) <= 0 or float(p.mediabox.height) <= 0 for p in document.pages):
                    raise ValueError("empty pages")
                details = {"pages": len(document.pages)}
            elif path.suffix.lower() == ".docx":
                from docx import Document
                document = Document(path)
                details = {"paragraphs": len(document.paragraphs), "tables": len(document.tables)}
            elif path.suffix.lower() == ".xlsx":
                from openpyxl import load_workbook
                document = load_workbook(path, read_only=True, data_only=False)
                details = {"sheets": len(document.sheetnames)}
                document.close()
            elif path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}:
                from PIL import Image
                with Image.open(path) as image:
                    image.verify()
                with Image.open(path) as image:
                    image.load()
                    details = {"width": image.width, "height": image.height}
            elif path.suffix.lower() == ".zip":
                with ZipFile(path) as archive:
                    if archive.testzip() is not None or len(set(archive.namelist())) != len(archive.namelist()):
                        raise ValueError("corrupt archive")
                    details = {"files": len(archive.namelist())}
            else:
                continue
        except Exception:
            raise ToolError("INVALID_OUTPUT", "生成文件无法重新打开，请重试") from None
        checks.append({"name": "output_reopen", "file": path.name, "status": "passed", **details})
    return checks
