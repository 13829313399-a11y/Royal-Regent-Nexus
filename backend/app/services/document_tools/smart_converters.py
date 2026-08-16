from __future__ import annotations

import hashlib
import re
from collections import OrderedDict
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader

from app.core.config import Settings
from app.schemas.document_studio import DocumentBlock, DocumentSnapshot
from app.services.document_studio.contracts import (
    DocumentBlockKind,
    DocumentExtractionRoute,
    DocumentExtractionSource,
)
from app.services.document_studio.evidence import extract_local_snapshot
from app.services.document_studio.extractors.qwen_ocr import cloud_ocr_page_numbers
from app.services.document_tools.contracts import (
    DocumentToolError,
    OcrResult,
    ProcessingMode,
    StructuredTable,
)
from app.services.document_tools.qwen_ocr import QwenOcrService
from app.services.document_tools.qwen_table_extractor import QwenTableExtractor
from app.services.document_tools.qwen_translation import QwenTranslationService
from app.services.pdf_to_excel import (
    PdfToExcelConversionError,
    PdfToExcelResult,
    convert_pdf_to_excel,
)
from app.services.pdf_to_word import (
    PdfToWordResult,
    convert_pdf_to_word,
    convert_pdf_to_word_layout,
)
from app.services.pdf_translation import PdfTranslationResult, convert_pdf_translation


@dataclass(frozen=True)
class SmartPdfToExcelResult(PdfToExcelResult):
    qwen_page_count: int = 0
    low_confidence_count: int = 0
    provider_model: str = ""
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class SmartPdfToWordResult(PdfToWordResult):
    qwen_page_count: int = 0
    provider_model: str = ""
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class SmartPdfTranslationResult(PdfTranslationResult):
    qwen_page_count: int = 0
    provider_model: str = ""
    warnings: tuple[str, ...] = ()


def _page_count(data: bytes, settings: Settings) -> int:
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        if reader.is_encrypted:
            raise ValueError("encrypted")
        count = len(reader.pages)
    except Exception as exc:
        raise DocumentToolError(
            "PDF_INVALID",
            "PDF 无法读取；文件可能已加密或损坏。",
            action="请移除 PDF 打开密码或重新导出文件。",
        ) from exc
    if count < 1:
        raise DocumentToolError("PDF_EMPTY", "PDF 中没有页面。")
    if count > settings.document_tool_max_pdf_pages:
        raise DocumentToolError(
            "PDF_PAGE_LIMIT_EXCEEDED",
            f"单次最多处理 {settings.document_tool_max_pdf_pages} 页 PDF。",
            action="请先拆分 PDF 后再处理。",
            status_code=413,
        )
    return count


def _local_snapshot(data: bytes) -> DocumentSnapshot:
    digest = hashlib.sha256(data).hexdigest()
    return extract_local_snapshot(
        data=data,
        source_artifact_id=f"aiart-{digest[:32]}",
        source_sha256=digest,
    )


def _selected_pages(
    snapshot: DocumentSnapshot,
    mode: ProcessingMode,
) -> tuple[int, ...]:
    if mode == ProcessingMode.LOCAL:
        return ()
    if mode == ProcessingMode.QWEN:
        return tuple(page.page_number for page in snapshot.pages)
    return cloud_ocr_page_numbers(snapshot)


def _safe_sheet_title(value: str, existing: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]+", "_", value).strip()[:31] or "千问结果"
    title = base
    suffix = 2
    while title in existing:
        marker = f"_{suffix}"
        title = f"{base[: 31 - len(marker)]}{marker}"
        suffix += 1
    existing.add(title)
    return title


def _blank_workbook(source_name: str, page_count: int) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "转换说明"
    sheet.append(["PDF 转 Excel", "转换结果"])
    sheet.append(["源文件", Path(source_name).name])
    sheet.append(["PDF 页数", str(page_count)])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=12)
        cell.fill = PatternFill("solid", fgColor="0F766E")
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 88
    return workbook


def _remove_qwen_page_sheets(workbook: Workbook, pages: tuple[int, ...]) -> None:
    for page_number in pages:
        prefix = f"第{page_number}页_"
        for worksheet in tuple(workbook.worksheets):
            if worksheet.title.startswith(prefix) and worksheet.title != "转换说明":
                workbook.remove(worksheet)


def _group_tables(
    page_tables: list[tuple[int, StructuredTable]],
) -> list[tuple[str, list[tuple[int, StructuredTable]]]]:
    grouped: OrderedDict[str, list[tuple[int, StructuredTable]]] = OrderedDict()
    for index, (page_number, table) in enumerate(page_tables, start=1):
        key = table.continuation_key.strip() or f"page-{page_number}-table-{index}"
        existing = grouped.get(key)
        if existing and tuple(cell.value for cell in existing[0][1].header) != tuple(
            cell.value for cell in table.header
        ):
            key = f"{key}-page-{page_number}-{index}"
        grouped.setdefault(key, []).append((page_number, table))
    return list(grouped.items())


def _append_qwen_tables(
    workbook: Workbook,
    *,
    tables: list[tuple[int, StructuredTable]],
    ocr: OcrResult,
) -> tuple[int, int]:
    existing = set(workbook.sheetnames)
    low_confidence = 0
    for table_index, (_key, group) in enumerate(_group_tables(tables), start=1):
        first_page, first = group[0]
        label = first.title.strip() or f"第{first_page}页_千问表格{table_index}"
        worksheet = workbook.create_sheet(_safe_sheet_title(label, existing))
        rows = [first.header]
        for _page_number, table in group:
            rows.extend(table.rows)
        for row_index, row in enumerate(rows, start=1):
            for column_index, source in enumerate(row, start=1):
                cell = worksheet.cell(
                    row=row_index, column=column_index, value=source.value
                )
                cell.number_format = "@"
                cell.quotePrefix = True
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="0F766E")
                elif source.confidence < 0.85:
                    low_confidence += 1
                    cell.fill = PatternFill("solid", fgColor="FFF3BF")
                    cell.comment = Comment(
                        f"千问识别置信度 {source.confidence:.0%}，建议与源 PDF 复核。",
                        "Royal Regent Nexus",
                    )
                worksheet.column_dimensions[cell.column_letter].width = max(
                    worksheet.column_dimensions[cell.column_letter].width or 10,
                    min(48, max(10, len(source.value) + 2)),
                )

    pages_with_tables = {page_number for page_number, _table in tables}
    for page in ocr.pages:
        if page.page_number in pages_with_tables or not page.blocks:
            continue
        worksheet = workbook.create_sheet(
            _safe_sheet_title(f"第{page.page_number}页_千问文字", existing)
        )
        worksheet.append(["识别文本", "置信度"])
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
        for block in page.blocks:
            worksheet.append([block.text, f"{block.confidence:.0%}"])
            worksheet.cell(worksheet.max_row, 1).number_format = "@"
            if block.confidence < 0.85:
                low_confidence += 1
                worksheet.cell(worksheet.max_row, 1).fill = PatternFill(
                    "solid", fgColor="FFF3BF"
                )
        worksheet.column_dimensions["A"].width = 88
        worksheet.column_dimensions["B"].width = 14
    return len(tables), low_confidence


def convert_pdf_to_excel_smart(
    pdf_bytes: bytes,
    source_file_name: str,
    *,
    settings: Settings,
    mode: ProcessingMode,
    ocr_service: QwenOcrService | None = None,
    table_extractor: QwenTableExtractor | None = None,
) -> SmartPdfToExcelResult:
    page_count = _page_count(pdf_bytes, settings)
    if mode == ProcessingMode.LOCAL:
        local = convert_pdf_to_excel(pdf_bytes, source_file_name)
        return SmartPdfToExcelResult(**local.__dict__)

    snapshot = _local_snapshot(pdf_bytes)
    selected = _selected_pages(snapshot, mode)
    if not selected:
        local = convert_pdf_to_excel(pdf_bytes, source_file_name)
        return SmartPdfToExcelResult(**local.__dict__)

    ocr_service = ocr_service or QwenOcrService(settings)
    table_extractor = table_extractor or QwenTableExtractor(settings)
    ocr = ocr_service.parse_pdf(pdf_bytes, page_numbers=selected)
    page_tables: list[tuple[int, StructuredTable]] = []
    for page in ocr.pages:
        extracted = table_extractor.extract_page(page)
        page_tables.extend((page.page_number, table) for table in extracted.tables)

    try:
        local = convert_pdf_to_excel(pdf_bytes, source_file_name)
        workbook = load_workbook(BytesIO(local.content))
    except PdfToExcelConversionError:
        local = None
        workbook = _blank_workbook(source_file_name, page_count)
    _remove_qwen_page_sheets(workbook, selected)
    qwen_table_count, low_confidence = _append_qwen_tables(
        workbook,
        tables=page_tables,
        ocr=ocr,
    )
    if len(workbook.worksheets) == 1:
        raise DocumentToolError(
            "DOCUMENT_NO_EXTRACTABLE_CONTENT",
            "未识别到可写入 Excel 的文字或表格。",
            action="请确认 PDF 不是空白页，并尝试上传更清晰的文件。",
        )
    summary = workbook["转换说明"]
    summary.append(["实际处理模式", mode.value])
    summary.append(["千问 OCR 页", str(len(selected))])
    summary.append(["千问结构化表格", str(qwen_table_count)])
    summary.append(["低置信单元格/内容块", str(low_confidence)])
    output = BytesIO()
    workbook.save(output)
    stem = (
        re.sub(r'[\\/:*?"<>|]+', "_", Path(source_file_name).stem).strip(" .")
        or "PDF文件"
    )
    return SmartPdfToExcelResult(
        content=output.getvalue(),
        page_count=page_count,
        table_count=(local.table_count if local else 0) + qwen_table_count,
        text_page_count=local.text_page_count if local else 0,
        ocr_page_count=(local.ocr_page_count if local else 0),
        output_file_name=f"{stem}_转换结果.xlsx",
        qwen_page_count=len(selected),
        low_confidence_count=low_confidence,
        provider_model=f"{settings.qwen_ocr_model}+{settings.qwen_table_model}",
    )


def _snapshot_with_qwen(
    snapshot: DocumentSnapshot,
    ocr: OcrResult,
) -> DocumentSnapshot:
    payload = snapshot.model_dump(mode="json")
    pages = {int(page["page_number"]): page for page in payload["pages"]}
    for parsed in ocr.pages:
        target = pages[parsed.page_number]
        width = float(target["width"])
        height = float(target["height"])
        target["blocks"] = []
        for index, block in enumerate(parsed.blocks, start=1):
            if block.bbox is None:
                top = max(0.0, (index - 1) * height / max(len(parsed.blocks), 1))
                bottom = min(height, index * height / max(len(parsed.blocks), 1))
                bbox = (0.0, top, width, bottom)
            else:
                left, top, right, bottom = block.bbox
                if max(block.bbox) <= 1:
                    bbox = (left * width, top * height, right * width, bottom * height)
                else:
                    bbox = (
                        max(0.0, min(left, width)),
                        max(0.0, min(top, height)),
                        max(0.0, min(right, width)),
                        max(0.0, min(bottom, height)),
                    )
            target["blocks"].append(
                DocumentBlock(
                    block_id=f"p{parsed.page_number}-b{index}",
                    kind=DocumentBlockKind.PARAGRAPH,
                    bbox=bbox,
                    raw_text=block.text,
                    normalized_text=block.text,
                    confidence=block.confidence,
                    source=DocumentExtractionSource.QWEN_OCR,
                    needs_review=block.confidence < 0.85,
                ).model_dump(mode="json")
            )
        target["extraction_route"] = DocumentExtractionRoute.QWEN_OCR.value
    return DocumentSnapshot.model_validate(payload)


def convert_pdf_to_word_smart(
    pdf_bytes: bytes,
    source_file_name: str,
    *,
    settings: Settings,
    mode: ProcessingMode,
    output_mode: str = "EDITABLE",
    ocr_service: QwenOcrService | None = None,
) -> SmartPdfToWordResult:
    _page_count(pdf_bytes, settings)
    if output_mode not in {"EDITABLE", "LAYOUT_PRESERVING"}:
        raise DocumentToolError(
            "PDF_TO_WORD_MODE_INVALID",
            "PDF 转 Word 输出模式无效。",
            action="请选择可编辑优先或版式保真。",
        )
    if output_mode == "LAYOUT_PRESERVING":
        if mode == ProcessingMode.QWEN:
            raise DocumentToolError(
                "QWEN_NOT_APPLICABLE_TO_LAYOUT",
                "版式保真输出不需要调用千问。",
                action="请选择自动或仅本地模式，或改用可编辑优先。",
            )
        local = convert_pdf_to_word_layout(pdf_bytes, source_file_name)
        return SmartPdfToWordResult(
            **local.__dict__,
            warnings=(
                "版式保真模式将每页作为清晰页面图像写入 Word，文字不可直接编辑。",
            ),
        )
    if mode == ProcessingMode.LOCAL:
        local = convert_pdf_to_word(pdf_bytes, source_file_name)
        return SmartPdfToWordResult(**local.__dict__)
    snapshot = _local_snapshot(pdf_bytes)
    selected = _selected_pages(snapshot, mode)
    if not selected:
        local = convert_pdf_to_word(pdf_bytes, source_file_name)
        return SmartPdfToWordResult(**local.__dict__)
    ocr = (ocr_service or QwenOcrService(settings)).parse_pdf(
        pdf_bytes,
        page_numbers=selected,
    )
    enhanced = _snapshot_with_qwen(snapshot, ocr)
    overrides = {
        page.page_number: tuple(
            (block.bbox[1], block.bbox[3], block.raw_text)
            for block in page.blocks
            if block.raw_text.strip()
        )
        for page in enhanced.pages
        if page.page_number in selected
    }
    result = convert_pdf_to_word(
        pdf_bytes,
        source_file_name,
        page_text_overrides=overrides,
    )
    return SmartPdfToWordResult(
        **result.__dict__,
        qwen_page_count=len(selected),
        provider_model=settings.qwen_ocr_model,
    )


def convert_pdf_translation_smart(
    pdf_bytes: bytes,
    source_file_name: str,
    *,
    settings: Settings,
    mode: ProcessingMode,
    requested_direction: str,
    layout: str,
    protected_tokens: tuple[str, ...],
    include_editable_docx: bool,
    terms: tuple[dict[str, str], ...] = (),
    translation_memory: tuple[dict[str, str], ...] = (),
    domain_prompt: str = "",
    ocr_service: QwenOcrService | None = None,
    translation_service: QwenTranslationService | None = None,
) -> SmartPdfTranslationResult:
    _page_count(pdf_bytes, settings)
    if mode == ProcessingMode.LOCAL:
        result = convert_pdf_translation(
            pdf_bytes,
            source_file_name,
            settings=settings,
            requested_direction=requested_direction,
            layout=layout,
            protected_tokens=protected_tokens,
            include_editable_docx=include_editable_docx,
        )
        return SmartPdfTranslationResult(**result.__dict__)

    snapshot = _local_snapshot(pdf_bytes)
    selected = cloud_ocr_page_numbers(snapshot)
    if selected:
        ocr = (ocr_service or QwenOcrService(settings)).parse_pdf(
            pdf_bytes,
            page_numbers=selected,
        )
        snapshot = _snapshot_with_qwen(snapshot, ocr)
    translator = translation_service or QwenTranslationService(settings)

    def translate_batch(values, direction):
        return translator.translate_batch(
            values,
            direction,
            terms=terms,
            translation_memory=translation_memory,
            domain_prompt=domain_prompt,
        )

    result = convert_pdf_translation(
        pdf_bytes,
        source_file_name,
        settings=settings,
        requested_direction=requested_direction,
        layout=layout,
        protected_tokens=protected_tokens,
        include_editable_docx=include_editable_docx,
        translator=translate_batch,
        snapshot_override=snapshot,
    )
    return SmartPdfTranslationResult(
        **result.__dict__,
        qwen_page_count=len(selected),
        provider_model=settings.qwen_translation_model,
    )
