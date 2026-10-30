"""Inline inspection journal; additive storage, no writes to Sales or reviewed history."""
import base64
import binascii
from contextlib import contextmanager
from datetime import date
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
import sqlite3
import time
from uuid import uuid4
import warnings

from fastapi import HTTPException
from PIL import Image, ImageOps, UnidentifiedImageError
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.services import qc_workflow as wf

MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
# Keep the template business fields; the final editing columns follow the requested result / remark / photos order.
COLUMNS = [('date', '日期'), ('customer', '客户'), ('poNo', 'PO#'), ('country', '出口国'),
           ('productNo', '货号'), ('productName', '产品名称'), ('quantity', '数量'),
           ('unitsPerCarton', '装箱'), ('cartonCount', '箱数'), ('production', '生产'),
           ('agency', '机构'), ('customerPo', '客PO#'), ('completion', '完成情况'),
           ('result', '验货结果'), ('remark', '备注'), ('photos', '照片')]
SCHEMA = '''
CREATE TABLE IF NOT EXISTS entries (
 id TEXT PRIMARY KEY, factory TEXT NOT NULL, source_key TEXT NOT NULL,
 base TEXT NOT NULL, result TEXT NOT NULL CHECK(result IN ('','PASS','FAIL','PENDING')),
 remark TEXT NOT NULL, revision INTEGER NOT NULL DEFAULT 0, edited INTEGER NOT NULL DEFAULT 0,
 updated_at TEXT NOT NULL, UNIQUE(factory,source_key));
CREATE TABLE IF NOT EXISTS photos (
 id TEXT PRIMARY KEY, entry_id TEXT NOT NULL REFERENCES entries(id), name TEXT NOT NULL,
 content BLOB NOT NULL, digest TEXT NOT NULL, origin TEXT, removed INTEGER NOT NULL DEFAULT 0,
 created_at TEXT NOT NULL, UNIQUE(entry_id,origin));
CREATE TABLE IF NOT EXISTS changes (
 id INTEGER PRIMARY KEY, entry_id TEXT NOT NULL REFERENCES entries(id), revision INTEGER NOT NULL,
 action TEXT NOT NULL, actor TEXT NOT NULL, created_at TEXT NOT NULL, snapshot TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS exports (
 id TEXT PRIMARY KEY, factory TEXT NOT NULL, actor TEXT NOT NULL, expires INTEGER NOT NULL, content BLOB NOT NULL);
'''


@contextmanager
def store(write=False):
    path = Path(settings.qc_department_data_dir).resolve() / '_inspection_sheet' / 'inspections.sqlite3'
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(path, timeout=20)
    conn.row_factory = sqlite3.Row
    try:
        conn.execute('PRAGMA foreign_keys=ON')
        conn.executescript(SCHEMA)
        migrate_pending(conn)
        conn.execute('BEGIN IMMEDIATE' if write else 'BEGIN')
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def scalar(value):
    return '' if value is None else str(value)


def base_fields(source):
    result = {key: source.get(key, '') for key, _ in COLUMNS if key not in {'date', 'remark', 'result', 'photos'}}
    start, end = wf.source_date(source.get('dateFrom')), wf.source_date(source.get('dateTo'))
    result.update(dateFrom=start, dateTo=end or start, sourceRevision=source.get('sourceRevision'),
                  sourceNote=scalar(source.get('sourceRemark', source.get('remark', ''))))
    result['actualDate'] = wf.source_date(source.get('actualDate'))
    for key in ('inspectionProblem','otherRemark','inspectionFactory','reportStatus','responsible','firstResponsible','secondResponsible'):
        if key in source:result[key]=source[key]
    result['referenceData']=bool(source.get('reference'))
    return result


def legacy_fields(factory, record, examples):
    base = base_fields({'customer':record.get('client'), 'poNo':record.get('orderNo'),
        'productNo':record.get('productNo'), 'productName':record.get('productName'),
        'quantity':record.get('qty'), 'remark':record.get('remark'),'actualDate':record.get('inspDate')})
    # Exact saved provenance + allowlisted full snapshot; never join by name or row position.
    match = re.search(r'^来源：[^\n]*；(schedule-example-row-\d+)\s*$', record.get('remark') or '', re.M)
    source = examples.get(match.group(1)) if match and wf.test_entry_allowed(factory, record) else None
    if source and all(scalar(record.get(a)).strip() == scalar(source.get(b)).strip()
                      for a,b in [('client','customer'),('orderNo','poNo'),('productNo','productNo')]):
        base = base_fields(source)
        base['testData'] = True
    base['actualDate']=wf.source_date(record.get('inspDate'))
    base['testData']=bool(base.get('testData') or wf.test_entry_allowed(factory,record))
    return base


def candidates(factory, records, tasks, sources, examples=(), references=()):
    """Use explicit source IDs. Distinct batches and duplicate SKUs remain independent."""
    example_map = {r['id']:r for r in examples}
    source_map = {r['id']:r for r in sources}
    linked = {t['legacyId'] for t in tasks if t.get('legacyId') is not None}
    result = []
    for record in records:
        if record['id'] in linked:
            continue
        base = legacy_fields(factory, record, example_map)
        base['legacyId'] = record['id']
        base['referenceData']=record['id'] in references
        result.append(('legacy:'+str(record['id']), base, record.get('result',''),
                       base['sourceNote'], []))
    for task in tasks:
        source = source_map.get(task['sourceId'], task['source'])
        base = base_fields(source)
        if task.get('legacyId') is not None:
            base = legacy_fields(factory, task.get('legacyOriginal') or {}, example_map)
            base['legacyId'] = task['legacyId']
        else:
            base['quantity'] = task.get('plannedQuantity') if task.get('plannedQuantity') is not None else base['quantity']
            if not base['dateFrom']:
                base['dateFrom'] = base['dateTo'] = wf.source_date(task.get('plannedDate'))
        base.update(workflowId=task['id'], batch=task['batch'], sourceId=task['sourceId'])
        execution = task['execution']
        base['actualDate']=wf.source_date(execution.get('actualDate'))
        base['referenceData']=bool(task.get('reference') or task.get('legacyId') in references)
        remark = execution.get('remark','')
        if base.get('testData') and remark == (task.get('legacyOriginal') or {}).get('remark'):
            remark = base['sourceNote']
        base['testData'] = bool(base.get('testData') or task.get('testEntry'))
        key = 'legacy:'+str(task['legacyId']) if task.get('legacyId') is not None else 'workflow:'+task['id']
        result.append((key,base,execution.get('result',''),remark or base['sourceNote'],task['attachments']))
    # Sales schedules enter this journal only through an explicitly created task.
    return result


def get(conn, factory, entry_id):
    row = conn.execute('SELECT * FROM entries WHERE id=? AND factory=?',(entry_id,factory)).fetchone()
    if row is None:
        raise HTTPException(404,'验货任务不存在')
    return row


def unpack(conn, row, include_photos=True):
    photos = [dict(f) for f in conn.execute('SELECT id,name FROM photos WHERE entry_id=? AND removed=0 ORDER BY created_at,id',(row['id'],))] if include_photos else []
    base=json.loads(row['base'])
    raw=base.get('sourceResult','')
    cumulative_result=raw if raw in {'REJ','COND','HOLD'} and not base.get('resultOverride') else row['result']
    return {**base, 'id':row['id'], 'sourceKey':row['source_key'], 'cumulativeResult':cumulative_result,
            'result':row['result'], 'remark':row['remark'], 'revision':row['revision'],
            'updatedAt':row['updated_at'], 'photos':photos}


def sync(factory, rows, copy_photos=True):
    ids = [];photo_warnings=set()
    with store(True) as conn:
        source_counts={}
        for _,base,_,_,_ in rows:
            if base.get('sourceId'):source_counts[base['sourceId']]=source_counts.get(base['sourceId'],0)+1
        for key,base,result,remark,attachments in rows:
            entry_id = sha256((factory+'\0'+key).encode()).hexdigest()[:32]
            previous = conn.execute('SELECT * FROM entries WHERE factory=? AND source_key=?',(factory,key)).fetchone()
            # A scheduled source becoming its first workflow keeps the same business identity.
            # Subsequent batches remain independent; never join on customer/PO/SKU.
            if not previous and base.get('workflowId'):
                linked=list(conn.execute("SELECT * FROM entries WHERE factory=? AND json_extract(base,'$.workflowId')=?",(factory,base['workflowId'])))
                source=conn.execute('SELECT * FROM entries WHERE factory=? AND source_key=?',(factory,'source:'+str(base.get('sourceId','')))).fetchone()
                if not linked and source:
                    if source_counts.get(base.get('sourceId'))==1:linked=[source]
                    else:base['lineageConflict']='原排期已拆为多个批次，无法自动归属先前填写的验货记录'
                if len(linked)==1:
                    previous=linked[0]
                    conn.execute('UPDATE entries SET source_key=? WHERE id=?',(key,previous['id']))
                elif len(linked)>1:base['lineageConflict']='同一流程关联了多个验货记录，需要核对'
            if previous:entry_id=previous['id']
            ids.append(entry_id)
            base['historicalResult'] = result if result not in ('',None,'PASS','FAIL','PENDING','REJ') else ''
            base['sourceResult']=result or ''
            base['resultOverride']=False
            if previous:
                old_base=json.loads(previous['base'])
                if 'resultOverride' in old_base:base['resultOverride']=old_base['resultOverride']
                else:
                    # Old entries already have immutable save receipts; photo/remark-only saves do not change judgment.
                    prior='FAIL' if result in ('FAIL','REJ') else result if result in ('PASS','PENDING') else ''
                    for change in conn.execute("SELECT snapshot FROM changes WHERE entry_id=? AND action='save' ORDER BY revision",(entry_id,)):
                        value=json.loads(change['snapshot']).get('result','')
                        if value!=prior:base['resultOverride']=True
                        prior=value
                stale=isinstance(base.get('sourceRevision'),int) and isinstance(old_base.get('sourceRevision'),int) and base['sourceRevision']<old_base['sourceRevision']
                if previous['base'] != wf.dump(base) and not stale:
                    conn.execute('UPDATE entries SET base=?,revision=revision+1 WHERE id=?',(wf.dump(base),entry_id))
            else:
                normalized = 'FAIL' if result in ('FAIL','REJ') else result if result in ('PASS','PENDING') else ''
                conn.execute('INSERT INTO entries VALUES(?,?,?,?,?,?,0,0,?)',
                    (entry_id,factory,key,wf.dump(base),normalized,remark or '',wf.now()))
            # Preserve existing photos as independent copies. Original reviewed evidence is immutable.
            for attachment in attachments if copy_photos else []:
                origin = base.get('workflowId','')+':'+attachment['id']
                if not attachment['mime'].startswith('image/') or conn.execute('SELECT 1 FROM photos WHERE entry_id=? AND origin=?',(entry_id,origin)).fetchone():
                    continue
                try:
                    content = wf.read_file(factory,base['workflowId'],attachment['id'])['content']
                    content = normalize_photo(content)
                except (HTTPException,OSError):
                    # Keep the original attachment untouched and retry on a later ordinary page refresh.
                    # An unavailable historical photo must not hide readable inspection text.
                    photo_warnings.add(entry_id);continue
                conn.execute('INSERT INTO photos VALUES(?,?,?,?,?,?,0,?)',
                    (uuid4().hex,entry_id,attachment['name'],content,sha256(content).hexdigest(),origin,wf.now()))
                conn.execute('UPDATE entries SET revision=revision+1 WHERE id=?',(entry_id,))
        # Preserve existing inspection history. Pre-push source caches stay hidden until
        # a confirmed task adopts them above, including any previously saved input.
        retained=[]
        for row in conn.execute('SELECT * FROM entries WHERE factory=?',(factory,)):
            if row['id'] in ids:
                item=unpack(conn,row,copy_photos)
                if row['id'] in photo_warnings:item['photoWarning']='来源照片暂不可用；刷新可重试，文字仍可保存和导入'
                retained.append(item)
            elif row['edited'] and not row['source_key'].startswith('source:'):
                item=unpack(conn,row,copy_photos)
                item['sourceWarning']='来源已移除、取消或生成了新批次；保留原验货记录，未自动合并。'
                retained.append(item)
        return {'items':retained, 'columns':[{'key':key,'label':label} for key,label in COLUMNS]}


def revision(row, body):
    if type(body.get('revision')) is not int or body['revision'] != row['revision']:
        raise HTTPException(409,'记录已更新，填写内容已保留；请重新读取服务器版本后重试')


def changed(conn, row, action, actor):
    conn.execute('UPDATE entries SET revision=revision+1,edited=1,updated_at=? WHERE id=?',(wf.now(),row['id']))
    current=get(conn,row['factory'],row['id'])
    item=unpack(conn,current)
    conn.execute('INSERT INTO changes(entry_id,revision,action,actor,created_at,snapshot) VALUES(?,?,?,?,?,?)',
        (row['id'],current['revision'],action,wf.dump(actor),wf.now(),wf.dump(item)))
    return item


def photo_revision_matches(conn,entry,expected):
    if type(expected) is not int or expected>entry['revision']:return False
    if expected==entry['revision']:return True
    changes=list(conn.execute('SELECT revision,action FROM changes WHERE entry_id=? AND revision>? AND revision<=? ORDER BY revision',
        (entry['id'],expected,entry['revision'])))
    # Ignore only a complete, audited sequence of photo-only revisions. Untracked source changes still conflict.
    return len(changes)==entry['revision']-expected and all(r['revision']==expected+i+1 and r['action'] in {'upload','remove-photo'} for i,r in enumerate(changes))

def save(factory, entry_id, body, actor, allow_photo_changes=False):
    if set(body) != {'revision','result','remark'} or body['result'] not in ('','PASS','FAIL','PENDING'):
        raise HTTPException(422,'验货结果只接受空值、PASS、FAIL 或 PENDING')
    remark=body['remark']
    if not isinstance(remark,str) or len(remark)>20000 or any(ord(c)<32 and c not in '\n\r\t' for c in remark):
        raise HTTPException(422,'备注最多 20000 字，不能含不可显示的控制字符')
    with store(True) as conn:
        row=get(conn,factory,entry_id)
        if not (allow_photo_changes and photo_revision_matches(conn,row,body['revision'])):revision(row,body)
        if body['result']!=row['result']:
            base=json.loads(row['base']);base['resultOverride']=True
            conn.execute('UPDATE entries SET base=? WHERE id=?',(wf.dump(base),entry_id))
        conn.execute('UPDATE entries SET result=?,remark=? WHERE id=?',(body['result'],remark,entry_id))
        return changed(conn,row,'save',actor)


def normalize_photo(content):
    if not content or len(content)>10*1024*1024:
        raise HTTPException(413,'单张照片须为 10 MB 以内的图片')
    try:
        try:
            from pillow_heif import register_heif_opener
            register_heif_opener()
        except ImportError:
            pass
        with warnings.catch_warnings():
            warnings.simplefilter('error', Image.DecompressionBombWarning)
            with Image.open(BytesIO(content)) as original:
                if original.format not in {'JPEG','PNG','WEBP','HEIF','HEIC'}:
                    raise ValueError()
                if original.width*original.height>40000000:
                    raise ValueError()
                original.load()
                photo=ImageOps.exif_transpose(original).convert('RGB')
                photo.thumbnail((2560,2560))
                output=BytesIO();photo.save(output,'JPEG',quality=90,optimize=True)
                return output.getvalue()
    except (UnidentifiedImageError,OSError,ValueError,Image.DecompressionBombError,Image.DecompressionBombWarning):
        raise HTTPException(422,'图片损坏或格式不支持，请选择 JPG、PNG、WebP，或将手机照片转为 JPEG 后重试') from None


def add_photo(factory, entry_id, body, actor):
    if set(body) != {'revision','name','content','requestKey'} or not re.fullmatch(r'[a-f0-9-]{16,64}',str(body['requestKey'])):
        raise HTTPException(422,'照片提交字段不正确')
    name=body['name']
    if not isinstance(name,str) or not 1<=len(name)<=255:
        raise HTTPException(422,'照片名称不正确')
    try:
        content=base64.b64decode(body['content'],validate=True)
    except (ValueError,TypeError,binascii.Error):
        raise HTTPException(422,'照片内容无效') from None
    content=normalize_photo(content)
    digest=sha256(content).hexdigest()
    with store(True) as conn:
        row=get(conn,factory,entry_id)
        origin='upload:'+body['requestKey']
        previous=conn.execute('SELECT digest FROM photos WHERE entry_id=? AND origin=?',(entry_id,origin)).fetchone()
        if previous:
            if previous['digest'] != digest:
                raise HTTPException(409,'同一上传编号不能用于不同照片')
            return unpack(conn,row)
        revision(row,body)
        count,total=conn.execute('SELECT count(*),coalesce(sum(length(content)),0) FROM photos WHERE entry_id=?',(entry_id,)).fetchone()
        if count>=50 or total+len(content)>100*1024*1024:
            raise HTTPException(413,'每条任务含历史照片最多 50 张、100 MB')
        conn.execute('INSERT INTO photos VALUES(?,?,?,?,?,?,0,?)',(uuid4().hex,entry_id,name,content,digest,origin,wf.now()))
        return changed(conn,row,'upload',actor)


def remove_photo(factory, entry_id, body, actor):
    if set(body) != {'revision','photoId'}:
        raise HTTPException(422,'删除字段不正确')
    with store(True) as conn:
        row=get(conn,factory,entry_id);revision(row,body)
        found=conn.execute('UPDATE photos SET removed=1 WHERE id=? AND entry_id=? AND removed=0',(body['photoId'],entry_id))
        if not found.rowcount:
            raise HTTPException(404,'照片不存在或已删除')
        return changed(conn,row,'remove-photo',actor)


def read_photo(factory, entry_id, photo_id):
    with store() as conn:
        get(conn,factory,entry_id)
        row=conn.execute('SELECT content FROM photos WHERE id=? AND entry_id=? AND removed=0',(photo_id,entry_id)).fetchone()
        if row is None:
            raise HTTPException(404,'照片不存在')
        return row['content']


def date_value(item):
    start,end=item.get('dateFrom',''),item.get('dateTo','')
    if start and (not end or start==end):
        return date.fromisoformat(start)
    return start+' 至 '+end if start else ''


def export_xlsx(factory, body):
    if set(body)!={'entries'} or not isinstance(body['entries'],list) or len(body['entries'])>10000:
        raise HTTPException(422,'导出范围不正确')
    requested=body['entries']
    if any(not isinstance(x,dict) or set(x)!={'id','revision'} or not isinstance(x['id'],str) for x in requested) or len({x['id'] for x in requested}) != len(requested):
        raise HTTPException(422,'导出任务编号重复或无效')
    items=[];photos={}
    with store() as conn:
        for wanted in requested:
            row=get(conn,factory,wanted['id']);revision(row,wanted);item=unpack(conn,row);items.append(item)
            photos[item['id']]=[dict(p) for p in conn.execute('SELECT content FROM photos WHERE entry_id=? AND removed=0 ORDER BY created_at,id',(item['id'],))]
    return export_records(items,photo_provider=lambda item:photos[item['id']])


def export_records(items,include_photos=True,photo_provider=None,epoch=None):
    """Shared original detail export layout; callers supply one saved ordered snapshot."""
    workbook=Workbook();sheet=workbook.active;sheet.title='验货明细'
    if epoch is not None:workbook.epoch=epoch
    columns=COLUMNS if include_photos else [(key,label) for key,label in COLUMNS if key!='photos']
    sheet.append([label for _,label in columns])
    widths=[23,18,27,24,29,33,12,12,12,15,15,22,22,14,48,36]
    for i,width in enumerate(widths[:len(columns)],1): sheet.column_dimensions[get_column_letter(i)].width=width
    line=Side(style='thin',color='CCD5D7')
    for cell in sheet[1]:
        cell.font=Font(name='Arial',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24584E')
        cell.alignment=Alignment(vertical='center');cell.border=Border(bottom=line)
    sheet.row_dimensions[1].height=26;sheet.freeze_panes='D2';sheet.auto_filter.ref=f'A1:{get_column_letter(len(columns))}{max(1,len(items)+1)}'
    total=0
    for index,item in enumerate(items,2):
        values=[date_value(item) if key=='date' else '' if key=='photos' else item.get(key,'') for key,_ in columns]
        for col,value in enumerate(values,1):
            cell=sheet.cell(index,col,value if value is not None else '')
            if isinstance(value,str):
                cell.data_type='s' # Never interpret PO/remark beginning with '=' as formulas.
                cell.number_format='@'
            if col==1 and isinstance(value,date):cell.number_format='m"月"d"日"'
            if col in (7,8,9) and value not in ('',None):
                try:
                    number=float(value)
                    if number!=number or abs(number)>1e12:raise ValueError()
                    cell.value=number;cell.number_format='#,##0.####'
                except (ValueError,TypeError):pass
            cell.font=Font(name='Arial',size=11);cell.alignment=Alignment(vertical='top',wrap_text=True)
            cell.border=Border(left=line,right=line,top=line,bottom=line)
        photos=photo_provider(item) if include_photos and photo_provider else []
        total+=sum(len(f['content']) for f in photos)
        if total>200*1024*1024:raise HTTPException(413,'照片总量超过 200 MB，请缩小筛选范围后导出')
        # Main row shows up to three photos; larger sets get an explicitly linked photo sheet.
        height=max(42,min(400,18*(1+len(item['remark'])//44+item['remark'].count('\n'))))
        for n,photo in enumerate(photos):
            image=ExcelImage(BytesIO(photo['content']));ratio=min(140/image.width,110/image.height)
            image.width*=ratio;image.height*=ratio
            from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
            from openpyxl.drawing.xdr import XDRPositiveSize2D
            from openpyxl.utils.units import pixels_to_EMU
            target=sheet;image_col=15;image_row=index-1;offset=n*150
            if len(photos)>3:
                if n==0:
                    target=workbook.create_sheet('照片-'+str(index-1))
                    target.append(['PO#','货号','任务编号','照片序号','照片'])
                    for col,width in [('A',30),('B',30),('C',36),('D',12),('E',35)]:target.column_dimensions[col].width=width
                    sheet.cell(index,16,'共 '+str(len(photos))+' 张；点击查看照片页')
                    sheet.cell(index,16).hyperlink="#'"+target.title+"'!A1"
                    target.freeze_panes='E2'
                target=workbook['照片-'+str(index-1)]
                target.append([scalar(item.get('poNo')),scalar(item.get('productNo')),item['id'],n+1,''])
                for cell in target[n+2]:
                    if isinstance(cell.value,str):cell.data_type='s';cell.number_format='@'
                target.row_dimensions[n+2].height=94
                image_col=4;image_row=n+1;offset=0
            image.anchor=OneCellAnchor(_from=AnchorMarker(col=image_col,row=image_row,colOff=pixels_to_EMU(6+offset),rowOff=pixels_to_EMU(6)),
                ext=XDRPositiveSize2D(pixels_to_EMU(image.width),pixels_to_EMU(image.height)))
            target.add_image(image)
        if photos:
            height=max(height,94)
            sheet.column_dimensions['P'].width=max(sheet.column_dimensions['P'].width,(min(len(photos),3)*150+12)/7)
        sheet.row_dimensions[index].height=height
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0;sheet.print_title_rows='1:1'
    output=BytesIO();workbook.save(output);return output.getvalue()


def prepare_export(factory, body, actor):
    content=export_xlsx(factory,body);stamp=int(time.time());token=uuid4().hex
    with store(True) as conn:
        conn.execute('DELETE FROM exports WHERE expires<?',(stamp,))
        total=conn.execute('SELECT coalesce(sum(length(content)),0) FROM exports').fetchone()[0]
        if total+len(content)>500*1024*1024:raise HTTPException(503,'导出缓存繁忙，请稍后重试')
        conn.execute('INSERT INTO exports VALUES(?,?,?,?,?)',(token,factory,actor,stamp+600,content))
    return {'url':'api/inspection-sheet/exports/'+token,'filename':'QC验货明细.xlsx'}


def download_export(factory, actor, token):
    with store() as conn:
        row=conn.execute('SELECT * FROM exports WHERE id=? AND factory=? AND actor=?',(token,factory,actor)).fetchone()
        if row is None:raise HTTPException(404,'下载不存在')
        if row['expires']<int(time.time()):raise HTTPException(410,'下载已过期，请重新导出')
        return row['content']


def migrate_pending(conn):
    """Transactional CHECK expansion; retain row IDs and all child evidence."""
    if 'PENDING' in conn.execute("SELECT sql FROM sqlite_master WHERE name='entries'").fetchone()[0]:return
    conn.execute('PRAGMA foreign_keys=OFF')
    try:
        conn.execute('BEGIN IMMEDIATE')
        if 'PENDING' not in conn.execute("SELECT sql FROM sqlite_master WHERE name='entries'").fetchone()[0]:
            ddl=SCHEMA.split(';')[0].replace('IF NOT EXISTS entries','entries_pending_upgrade')
            conn.execute(ddl)
            conn.execute('INSERT INTO entries_pending_upgrade SELECT * FROM entries')
            conn.execute('DROP TABLE entries')
            conn.execute('ALTER TABLE entries_pending_upgrade RENAME TO entries')
            if conn.execute('PRAGMA foreign_key_check').fetchone():raise RuntimeError('Pending migration foreign-key validation failed')
        conn.commit()
    except Exception:conn.rollback();raise
    finally:conn.execute('PRAGMA foreign_keys=ON')
