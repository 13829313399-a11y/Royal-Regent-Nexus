# -*- coding: utf-8 -*-
"""Spin Master schedule matching, updating, and new-order export."""
from __future__ import annotations

import os
import re
import shutil
from collections import Counter, defaultdict
from copy import copy
from datetime import datetime
from typing import Any

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter


APP_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_MASTER_PATH = os.path.join(APP_DIR, "source_files", "华登-Spin客", "2026年Spin Master排货表(中国）.xlsx")

MASTER_SHEET = "2026年未验货订单"
SEARCH_SHEETS = ("2026年未验货订单", "2026已验未走货订单", "2026已走货订单 ")

BLUE_FILL = PatternFill(start_color="00B0F0", end_color="00B0F0", fill_type="solid")
HKD_RATE = 7.75

COL = {
    "po_date": 2,
    "change_note": 3,
    "contact": 4,
    "customer_po": 5,
    "contract": 6,
    "customer": 7,
    "version": 8,
    "item": 9,
    "cn_name": 10,
    "qty": 11,
    "inner": 12,
    "outer": 13,
    "total_box": 14,
    "special_note": 15,
    "carton_mark": 16,
    "customer_label": 17,
    "customer_material": 18,
    "material_reply": 19,
    "plastic_reply": 20,
    "carton_reply": 21,
    "injection_reply": 22,
    "line_start": 23,
    "complete_date": 24,
    "date_code": 25,
    "workshop_recheck": 26,
    "inspection_date": 27,
    "ship_date": 28,
    "booking_date": 29,
    "delivery_place": 30,
    "unit_hkd": 31,
    "unit_usd": 32,
    "total_hkd": 33,
    "total_usd": 34,
    "assembly_quote_hkd": 35,
    "english_name": 36,
    "pre_customer_material_date": 37,
    "certificate": 38,
    "carton_l": 39,
    "carton_w": 40,
    "carton_h": 41,
    "carton_cbm": 42,
    "total_cbm": 43,
    "gross_weight": 44,
    "net_weight": 45,
    "ship_country": 46,
    "so_no": 47,
    "container_type": 48,
    "warehouse_date": 49,
    "truck_detail": 50,
    "brand": 51,
    "system_date": 52,
}

EXPORT_END_COL = 188
CORE_END_COL = 30


def excel_serial(dt: datetime | None) -> int | None:
    if not dt:
        return None
    return (dt - datetime(1899, 12, 30)).days


def parse_date(value: Any) -> datetime | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value
    text = str(value)[:10]
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def normalize_contract(value: Any) -> str:
    text = str(value or "").strip()
    text = re.sub(r"\s+", "", text)
    if text.endswith(".0"):
        text = text[:-2]
    return text


def item_key(value: Any) -> str:
    text = str(value or "").upper().replace(" ", "")
    numbers = re.findall(r"\d{4,}", text)
    if len(numbers) >= 3:
        return "/".join(numbers[:3])
    return "/".join(numbers)


def line_contract(po_number: Any, line_no: Any) -> str:
    po = normalize_contract(po_number)
    suffix = normalize_contract(line_no)
    return f"{po}/{suffix}" if po and suffix else po


def _display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = str(value)
    return text[:-2] if text.endswith(".0") else text


def _value_equal(old: Any, new: Any) -> bool:
    if old is None and new in (None, ""):
        return True
    if isinstance(old, str) and old.strip().startswith("="):
        return True
    if isinstance(old, datetime) and isinstance(new, int):
        return excel_serial(old) == new
    try:
        if old not in (None, "") and new not in (None, ""):
            # 1分钱内视为相同：消除人工全精度乘积与单据两位小数的尾数噪音。
            return abs(float(old) - float(new)) < 0.011
    except (TypeError, ValueError):
        pass
    return _display(old).strip() == _display(new).strip()


class MasterIndex:
    def __init__(self) -> None:
        # 合同号 -> [(表名, 行号, 该行数量), ...]；同合同可能分批出货占多行，
        # 匹配时按数量配对，绝不能单行映射（后行会覆盖前行）。
        self.index: dict[str, list[tuple[str, int, Any]]] = {}
        self.catalog_by_item: dict[str, dict] = {}
        self.catalog_by_customer_item: dict[tuple[str, str], dict] = {}
        self.headers: dict[int, str] = {}
        self.max_row = 1
        self.inheritance_conflicts: list[dict[str, Any]] = []


def build_index(filepath: str) -> MasterIndex:
    idx = MasterIndex()
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    try:
        catalog_counters = defaultdict(lambda: defaultdict(Counter))
        catalog_samples = defaultdict(dict)
        customer_counters = defaultdict(lambda: defaultdict(Counter))
        customer_samples = defaultdict(dict)

        if MASTER_SHEET in wb.sheetnames:
            ws = wb[MASTER_SHEET]
            idx.max_row = ws.max_row
            idx.headers = {c: str(ws.cell(1, c).value or "") for c in range(1, min(EXPORT_END_COL, ws.max_column) + 1)}

        for sheet in SEARCH_SHEETS:
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            for r, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=min(EXPORT_END_COL, ws.max_column)), start=2):
                contract = normalize_contract(row[COL["contract"] - 1].value)
                key = item_key(row[COL["item"] - 1].value)
                customer = str(row[COL["customer"] - 1].value or "").strip().upper()
                if contract:
                    idx.index.setdefault(contract, []).append((sheet, r, row[COL["qty"] - 1].value))
                if not key:
                    continue
                for name, col in _CATALOG_COLS.items():
                    if col > len(row):
                        continue
                    val = row[col - 1].value
                    if val in (None, ""):
                        continue
                    catalog_counters[key][name].update([val])
                    catalog_samples[key].setdefault(name, val)
                    if customer:
                        customer_counters[(customer, key)][name].update([val])
                        customer_samples[(customer, key)].setdefault(name, val)

        for key in catalog_counters:
            idx.catalog_by_item[key] = _counter_to_catalog(
                catalog_counters[key], catalog_samples[key], idx.inheritance_conflicts, "全部客户", key
            )
        for key in customer_counters:
            idx.catalog_by_customer_item[key] = _counter_to_catalog(
                customer_counters[key], customer_samples[key], idx.inheritance_conflicts, key[0], key[1]
            )
        return idx
    finally:
        wb.close()


_CATALOG_COLS = {
    "contact": COL["contact"],
    "customer": COL["customer"],
    "version": COL["version"],
    "item_value": COL["item"],
    "cn_name": COL["cn_name"],
    "english_name": COL["english_name"],
    "inner": COL["inner"],
    "outer": COL["outer"],
    "special_note": COL["special_note"],
    "carton_mark": COL["carton_mark"],
    "customer_label": COL["customer_label"],
    "customer_material": COL["customer_material"],
    "material_reply": COL["material_reply"],
    "plastic_reply": COL["plastic_reply"],
    "carton_reply": COL["carton_reply"],
    "injection_reply": COL["injection_reply"],
    "delivery_place": COL["delivery_place"],
    "carton_l": COL["carton_l"],
    "carton_w": COL["carton_w"],
    "carton_h": COL["carton_h"],
    "carton_cbm": COL["carton_cbm"],
    "gross_weight": COL["gross_weight"],
    "net_weight": COL["net_weight"],
    "ship_country": COL["ship_country"],
    "brand": COL["brand"],
}


_PRODUCT_NAME_FIELDS = {"cn_name", "english_name"}


def _name_key(value: Any) -> str:
    return re.sub(r"[\W_]+", "", str(value or "").strip().casefold())


def _counter_to_catalog(
    counter_map: dict[str, Counter],
    samples: dict,
    conflicts: list[dict[str, Any]],
    scope: str,
    item: str,
) -> dict:
    result = {}
    for key, counter in counter_map.items():
        if key in _PRODUCT_NAME_FIELDS:
            values = {}
            for value in counter:
                normalized = _name_key(value)
                if normalized:
                    values.setdefault(normalized, value)
            if len(values) == 1:
                result[key] = next(iter(values.values()))
            elif len(values) > 1:
                conflicts.append({
                    "scope": scope,
                    "item": item,
                    "field": key,
                    "values": [str(value) for value in list(values.values())[:8]],
                })
            continue
        try:
            result[key] = counter.most_common(1)[0][0]
        except TypeError:
            result[key] = samples.get(key)
    return result


def _catalog_for(idx: MasterIndex, customer: str, key: str) -> dict:
    # 产品名称可以按货号全局继承；其余包装、联系人、标签等只能取
    # 同一客户+货号，避免同货号在不同客户之间串资料。
    global_catalog = idx.catalog_by_item.get(key, {})
    cat = {
        name: global_catalog[name]
        for name in _PRODUCT_NAME_FIELDS
        if global_catalog.get(name) not in (None, "")
    }
    cat.update(idx.catalog_by_customer_item.get((str(customer or "").strip().upper(), key), {}))
    return cat


def _round_money(value: float) -> float:
    return round(float(value or 0), 6)


def _compose_row(order: dict, line: dict, idx: MasterIndex) -> dict[int, Any]:
    po_date = parse_date(order.get("po_date"))
    ship_date = parse_date(line.get("ship_date"))
    contract = line_contract(order.get("po_number"), line.get("line_no"))
    key = line.get("item_key") or "/".join(
        part for part in (line.get("material_group"), line.get("material_number"), line.get("sales_material")) if part
    )
    cat = _catalog_for(idx, line.get("customer") or "", key)

    unit_usd = _round_money(line.get("unit_price") or 0)
    total_usd = _round_money((line.get("qty") or 0) * unit_usd)
    unit_hkd = _round_money(unit_usd * HKD_RATE) if unit_usd else ""
    total_hkd = _round_money(total_usd * HKD_RATE) if total_usd else ""
    outer = _outer_from_description(line.get("description_en", "")) or cat.get("outer", "")

    values = {
        COL["po_date"]: excel_serial(po_date),
        COL["contact"]: order.get("contact") or cat.get("contact", ""),
        COL["customer_po"]: line.get("customer_po", ""),
        COL["contract"]: contract,
        COL["customer"]: cat.get("customer") or line.get("customer", "") or order.get("customer", ""),
        COL["version"]: cat.get("version") or _version_from_description(line.get("description_en", "")),
        COL["item"]: _format_item(line, cat),
        COL["cn_name"]: cat.get("cn_name", ""),
        COL["qty"]: int(line.get("qty")) if float(line.get("qty") or 0).is_integer() else line.get("qty", ""),
        COL["inner"]: cat.get("inner", ""),
        COL["outer"]: outer,
        COL["special_note"]: cat.get("special_note", ""),
        COL["carton_mark"]: cat.get("carton_mark", ""),
        COL["customer_label"]: cat.get("customer_label", ""),
        COL["customer_material"]: cat.get("customer_material", ""),
        COL["material_reply"]: cat.get("material_reply", ""),
        COL["plastic_reply"]: cat.get("plastic_reply", ""),
        COL["carton_reply"]: cat.get("carton_reply", ""),
        COL["injection_reply"]: cat.get("injection_reply", ""),
        # SPIN 总汇规则表把验货期、预计订仓列为人工字段，PO 没有来源时留空。
        COL["inspection_date"]: "",
        COL["ship_date"]: excel_serial(ship_date),
        COL["booking_date"]: "",
        COL["delivery_place"]: cat.get("delivery_place", ""),
        COL["unit_hkd"]: unit_hkd,
        COL["unit_usd"]: unit_usd,
        COL["total_hkd"]: total_hkd,
        COL["total_usd"]: total_usd,
        COL["english_name"]: cat.get("english_name") or line.get("description_en", ""),
        COL["carton_l"]: cat.get("carton_l", ""),
        COL["carton_w"]: cat.get("carton_w", ""),
        COL["carton_h"]: cat.get("carton_h", ""),
        COL["carton_cbm"]: cat.get("carton_cbm", ""),
        COL["gross_weight"]: cat.get("gross_weight", ""),
        COL["net_weight"]: cat.get("net_weight", ""),
        COL["ship_country"]: cat.get("ship_country", ""),
        COL["brand"]: cat.get("brand", "SPINMASTER"),
    }
    return values


def _version_from_description(description: str) -> str:
    m = re.search(r"\b([A-Z]{3,5})\d*pk", description or "", re.I)
    return f"{m.group(1).upper()}版本" if m else ""


def _outer_from_description(description: str) -> int | str:
    """Spin rule: GML4pkSLD means 4 pieces per outer carton."""
    match = re.search(r"\bGML\s*(\d+)\s*pkSLD\b", description or "", re.I)
    return int(match.group(1)) if match else ""


def _format_item(line: dict, cat: dict) -> str:
    key = "/".join(
        str(line.get(name) or "").strip()
        for name in ("material_group", "material_number", "sales_material")
        if line.get(name)
    )
    sales = str(line.get("sales_order") or "").strip()
    item = str(line.get("line_item") or "").strip()
    tail = f"{sales}-{item}" if sales and item else ""
    if key and tail:
        return f"{key}/{tail}"
    return key or cat.get("item_value", "")


def write_orders(master_path: str, orders: list[dict], export_dir: str) -> dict:
    if not os.path.exists(master_path):
        return {"ok": False, "msg": f"排期文件不存在: {master_path}"}

    idx = build_index(master_path)
    new_rows = []
    new_details = []
    database_rows = []
    warnings = []
    inherited_rows = 0
    for order in orders:
        if not normalize_contract(order.get("po_number")):
            warnings.append(f"{order.get('filename', '')}: 未识别到 PO 号码，请在导出表中人工复核")
        for line in order.get("lines", []):
            values = _compose_row(order, line, idx)
            new_rows.append(values)
            new_details.append(values.get(COL["contract"], "") or values.get(COL["item"], ""))
            database_rows.append({
                "order_date": order.get("po_date", ""),
                "po_no": line.get("customer_po", ""),
                "contract_no": order.get("po_number", ""),
                "production_no": line.get("sales_order", ""),
                "client_name": line.get("customer") or order.get("customer", ""),
                "item_no": values.get(COL["item"], "") or line.get("item_key", ""),
                "product_name": values.get(COL["cn_name"], "") or line.get("description_en", ""),
                "quantity": line.get("qty"),
                "unit": line.get("unit", ""),
                "inspection_date": values.get(COL["inspection_date"], ""),
                "ship_date": line.get("ship_date", ""),
            })
            if values.get(COL["cn_name"]) or values.get(COL["english_name"]):
                inherited_rows += 1
            if not values.get(COL["outer"]):
                warnings.append(
                    f"{order.get('filename', '')} / {values.get(COL['item'], '')}: "
                    "未从 GMLxxpkSLD 或同客户排期识别到外箱装箱数，总箱数留空"
                )

    inheritance = {
        "rule": "same_schedule_same_item_unique_name",
        "mapped_items": sum(
            1
            for catalog in idx.catalog_by_item.values()
            if catalog.get("cn_name") or catalog.get("english_name")
        ),
        "applied_rows": inherited_rows,
        "conflicts": idx.inheritance_conflicts,
    }
    if idx.inheritance_conflicts:
        warnings.append(
            f"排期中有 {len(idx.inheritance_conflicts)} 组货号/品名冲突，相关品名未自动继承，请查看“排期继承检查”"
        )
    export_file = (
        generate_new_rows_excel(new_rows, master_path, export_dir, inheritance)
        if new_rows else ""
    )
    return {
        "ok": True,
        "msg": f"已按新单生成 {len(new_rows)} 行 Excel，原排期未修改" if new_rows else "没有识别到可生成的新单行",
        "modified": 0,
        "new_count": len(new_rows),
        "mod_details": [],
        "new_details": new_details,
        "rows": database_rows,
        "export_file": export_file,
        "warnings": warnings,
        "inheritance": inheritance,
        "mode": "new_order_only",
    }


_UPDATE_COLS = {
    # 来单期不覆盖：业务口径是"客户确认下单的邮件日期"，PO 日期只作新单参考。
    COL["contact"]: "联系人",
    COL["customer_po"]: "客PO",
    COL["contract"]: "合同号",
    COL["qty"]: "数量",
    COL["outer"]: "外箱",
    COL["inspection_date"]: "验货期",
    COL["ship_date"]: "PO走货期",
    COL["booking_date"]: "预计订仓时间",
    COL["unit_hkd"]: "单价/港币",
    COL["unit_usd"]: "单价/美金",
    COL["total_hkd"]: "总货价/港币",
    COL["total_usd"]: "总货价/美金",
    COL["english_name"]: "英文品名",
}


def _update_existing_row(ws, row_no: int, values: dict[int, Any]) -> list[str]:
    changes = []
    qty_or_outer_changed = False
    for col, label in _UPDATE_COLS.items():
        new = values.get(col)
        if new in (None, ""):
            continue
        cell = ws.cell(row_no, col)
        old = cell.value
        if _value_equal(old, new):
            continue
        cell.value = new
        cell.fill = BLUE_FILL
        changes.append(f"{label} {_display(old)} -> {_display(new)}")
        if col in (COL["qty"], COL["outer"]):
            qty_or_outer_changed = True

    if qty_or_outer_changed:
        total_cell = ws.cell(row_no, COL["total_box"])
        formula = f'=IF(M{row_no}=0,"",K{row_no}/M{row_no})'
        if total_cell.value != formula:
            total_cell.value = formula
            total_cell.fill = BLUE_FILL
            changes.append("总箱数 公式")
    return changes


def generate_new_rows_excel(
    rows: list[dict[int, Any]],
    master_path: str,
    output_dir: str,
    inheritance: dict[str, Any] | None = None,
) -> str:
    src = openpyxl.load_workbook(master_path)
    try:
        src_ws = src[MASTER_SHEET]
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "新单数据"
        _copy_header(src_ws, ws, EXPORT_END_COL)
        _write_rows(src_ws, ws, rows, EXPORT_END_COL)

        core = wb.create_sheet("仓单PO模板")
        _copy_header(src_ws, core, CORE_END_COL)
        _write_rows(src_ws, core, rows, CORE_END_COL)

        _autosize(ws, EXPORT_END_COL)
        _autosize(core, CORE_END_COL)
        if inheritance and (inheritance.get("applied_rows") or inheritance.get("conflicts")):
            check = wb.create_sheet("排期继承检查")
            headers = ("状态", "范围", "货号", "字段", "排期值")
            for col, title in enumerate(headers, 1):
                check.cell(1, col, title).fill = BLUE_FILL
            check.cell(2, 1, "已继承标准品名")
            check.cell(2, 2, "新单")
            check.cell(2, 3, f"{inheritance.get('applied_rows', 0)} 行")
            row_no = 3
            for conflict in inheritance.get("conflicts", []):
                values = (
                    "排期存在冲突，未自动继承",
                    conflict.get("scope", ""),
                    conflict.get("item", ""),
                    conflict.get("field", ""),
                    " / ".join(conflict.get("values", [])),
                )
                for col, value in enumerate(values, 1):
                    check.cell(row_no, col, value)
                row_no += 1
            check.freeze_panes = "A2"
            for col, width in enumerate((26, 18, 20, 18, 60), 1):
                check.column_dimensions[get_column_letter(col)].width = width
        os.makedirs(output_dir, exist_ok=True)
        fname = f"SpinMaster新单_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.xlsx"
        out = os.path.join(output_dir, fname)
        wb.save(out)
        wb.close()
        return fname
    finally:
        src.close()


def _copy_header(src_ws, dest_ws, end_col: int) -> None:
    for col in range(1, end_col + 1):
        src_cell = src_ws.cell(1, col)
        dest = dest_ws.cell(1, col, src_cell.value)
        if src_cell.has_style:
            dest.font = copy(src_cell.font)
            dest.fill = copy(src_cell.fill)
            dest.border = copy(src_cell.border)
            dest.alignment = copy(src_cell.alignment)
            dest.number_format = src_cell.number_format
        width = src_ws.column_dimensions[get_column_letter(col)].width
        if width:
            dest_ws.column_dimensions[get_column_letter(col)].width = width
    dest_ws.freeze_panes = "B2"


def _write_rows(src_ws, ws, rows: list[dict[int, Any]], end_col: int) -> None:
    style_row = 2 if src_ws.max_row >= 2 else 1
    for offset, data in enumerate(rows, start=2):
        if src_ws.row_dimensions[style_row].height:
            ws.row_dimensions[offset].height = src_ws.row_dimensions[style_row].height
        for col in range(1, end_col + 1):
            value = data.get(col, "")
            if col == COL["total_box"] and data.get(COL["qty"]) not in (None, ""):
                value = f'=IF(M{offset}=0,"",K{offset}/M{offset})'
            cell = ws.cell(offset, col, value)
            source = src_ws.cell(style_row, col)
            if source.has_style:
                cell.font = copy(source.font)
                cell.fill = copy(source.fill)
                cell.border = copy(source.border)
                cell.alignment = copy(source.alignment)
                cell.number_format = source.number_format
                cell.protection = copy(source.protection)
            if col in (COL["po_date"], COL["inspection_date"], COL["ship_date"]) and value not in (None, ""):
                cell.number_format = "yyyy/m/d"
            if col in (COL["customer_po"], COL["contract"]):
                cell.number_format = "@"


def _autosize(ws, end_col: int) -> None:
    for col in range(1, end_col + 1):
        letter = get_column_letter(col)
        if ws.column_dimensions[letter].width:
            continue
        ws.column_dimensions[letter].width = 12
