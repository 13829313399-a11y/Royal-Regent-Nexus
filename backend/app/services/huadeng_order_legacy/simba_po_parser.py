from __future__ import annotations

import io
import re
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook
from .common_pdf_text import extract_pdf_text

from .simba_schedule import add_derived_fields, read_schedule


DATE_FORMATS = ("%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d", "%Y/%m/%d", "%d.%m.%Y")


def parse_date_text(value: str | None) -> str | None:
    if not value:
        return None
    text = value.strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def read_pdf_text(source: str | Path | bytes | BinaryIO) -> tuple[str, int]:
    text, pages, _used_ocr = extract_pdf_text(source)
    return text, pages


def search(pattern: str, text: str, flags: int = re.IGNORECASE) -> str | None:
    match = re.search(pattern, text, flags)
    if not match:
        return None
    return next((g.strip() for g in match.groups() if g is not None), None)


def parse_reference_line(text: str) -> tuple[str | None, str | None, str | None]:
    line = search(r"Your\s+Reference\s+(.+)", text)
    if not line:
        return None, None, None
    bracket_customer = search(r"\[\s*([^\]]+?)\s*\]", line)
    without_bracket = re.sub(r"\[[^\]]+\]", " ", line)
    tokens = [token for token in re.split(r"\s+", without_bracket) if token]
    customer_po = None
    for token in tokens:
        if re.search(r"\d", token) and not token.upper() in {"DAP", "FOB", "LCL", "FCL"}:
            customer_po = token
            break
    trade_term = next((token for token in tokens if token.upper() in {"DAP", "FOB", "LCL", "FCL"}), None)
    return bracket_customer, customer_po, trade_term


def parse_line_items(text: str) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    line_pattern = re.compile(
        r"^\s*(\d+)\s+([A-Za-z0-9][A-Za-z0-9./()_-]*)\s+(.+?)\s+(\d{6,})\s+"
        r"([\d,]+(?:\.\d+)?)\s+(Each|PCS|PC|SET|Set|Unit|EA)\s+"
        r"([\d,]+(?:\.\d+)?)\s+([\d,]+(?:\.\d+)?)\s*$",
        re.IGNORECASE,
    )
    for raw_line in text.splitlines():
        line = re.sub(r"\s+", " ", raw_line).strip()
        match = line_pattern.match(line)
        if not match:
            continue
        _, item_no, description, _, quantity, _, unit_price, net_value = match.groups()
        if item_no.upper() in {"LCL", "FCL"}:
            continue
        records.append(
            {
                "item_no": item_no,
                "product_name": description,
                "quantity": float(quantity.replace(",", "")),
                "unit_price_usd": float(unit_price.replace(",", "")),
                "total_usd": float(net_value.replace(",", "")),
            }
        )
    return records


def _release_num(value: str) -> float:
    text = str(value or "").strip()
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        head, _, tail = text.rpartition(",")
        text = head.replace(",", "") + "." + tail if len(tail) == 2 and head else text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return 0.0


def _release_date(value: str | None) -> str | None:
    if not value:
        return None
    try:
        return datetime.strptime(value.strip(), "%d.%b.%Y").date().isoformat()
    except ValueError:
        return value


def workbook_text(source: str | Path | bytes | BinaryIO) -> str:
    """把整个工作簿按行拼成文本（转换版 PO 的兜底解析输入）。"""
    if isinstance(source, bytes):
        stream: Any = io.BytesIO(source)
    elif isinstance(source, (str, Path)):
        stream = io.BytesIO(Path(source).read_bytes())
    else:
        stream = source
    wb = load_workbook(stream, data_only=True, read_only=True)
    parts: list[str] = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None and str(v).strip()]
                if vals:
                    parts.append("  ".join(vals))
    finally:
        wb.close()
    return "\n".join(parts)


# Simba Dickie Release Order 数据行，兼容三种转换变体：
#   "70,560 PC  500054135/ 10  300480918/ 10  4.70  HKD  14.MAY.2026"（单行）
#   "120 PC  500047545/10  300471923/10  61.60·HKD  31.JAN.2026"（中点分隔币种）
#   "264 PC  500055145/10  300493420/10  Unit Price Delivery Date\n40.45 HKD 09.JUN.2026"（拆两行）
RELEASE_LINE_RE = re.compile(
    r"([\d.,]+)\s*PC\s+(\d{6,10})\s*/\s*(\d+)\s+(\d{6,10})\s*/\s*(\d+)"
    r"[\s\S]{0,80}?([\d.,]+)\s*[·\s]\s*([A-Z]{3})\s+(\d{1,2}\.[A-Za-z]{3,4}\.\d{4})",
    re.IGNORECASE,
)


PORT_COUNTRY_MAP = {
    "GENOA": "意大利",
    "GDANSK": "波兰",
    "HAMBURG": "德国",
    "ISTANBUL": "土耳其",
    "MELBOURNE": "澳大利亚",
    "ROTTERDAM": "荷兰",
    "SOUTHAMPTON": "英国",
}


def _release_reference(text: str) -> tuple[str | None, str | None]:
    """Return the blue-guide Reference value without the OCR-only SC prefix."""
    # Scanned pages often OCR the leading ``S`` in ``SC`` as ``$`` or ``5``.
    match = re.search(r"Reference:\s*(?:[S$5]C)?(\d{9})\s*/\s*(\d{1,4})", text, re.I)
    if not match:
        # WPS occasionally recognises "/" as "1", e.g. SC70013997711000.
        match = re.search(r"Reference:\s*(?:SC)?(\d{9})(\d{2,5})\b", text, re.I)
        if not match:
            return None, None
        line = match.group(2)
        if len(line) >= 2 and line.startswith("1"):
            line = line[1:]
    else:
        line = match.group(2)
    return f"{match.group(1)}/{line}", line


def _release_contact(text: str) -> str | None:
    match = re.search(r"Pers\.\s*respons\.?\s*[:.]?\s*([^\n]+)", text, re.I)
    if not match:
        return None
    contact = re.split(r"\s*-\s*Tel\s*:|\s+Tel\s*:", match.group(1), maxsplit=1, flags=re.I)[0]
    contact = re.sub(r"\s+", " ", contact).strip(" :-")
    return contact or None


def _release_remarks(text: str) -> str | None:
    matches = list(re.finditer(r"\bRemarks\s*:\s*", text, re.I))
    if not matches:
        return None
    # The actual PO requirements normally appear on the last page/last Remarks block.
    tail = text[matches[-1].end():]
    tail = re.split(
        r"\bPort\s+of\s+discharge\s*:|\*+\s*This\s+is\s+a?\s*final\s+order",
        tail,
        maxsplit=1,
        flags=re.I,
    )[0]
    value = re.sub(r"\s+", " ", tail).strip(" *:-")
    return value[:4000] or None


def _release_customer(remarks: str | None, text: str) -> str | None:
    match = re.search(r"\bCustomer\s*:\s*([^\n]{2,160})", remarks or "", re.I)
    if not match:
        match = re.search(r"\bCustomer\s*:\s*([^\n]{2,160})", text, re.I)
    if match:
        customer = match.group(1)
        customer = re.split(
            r"\s*(?:-+>|>|REMARKS\s*:|\d+\)|[\u4e00-\u9fff])",
            customer,
            maxsplit=1,
            flags=re.I,
        )[0]
        customer = re.sub(r"\s+", " ", customer).strip(" ,;:-")
        if customer:
            return customer
    match = None
    for source in (remarks or "", text):
        match = re.search(
            r"(?:This\s+)?orders?\s+(?:is\s+)?for\s+(?:cust(?:omer)?\s*)?"
            r"((?:Simba|ST)\s*(?:Toys\s*)?"
            r"(?:Germany|Italy|Turkey|Hungary|Poland|France|Australia|UK))",
            source,
            re.I,
        )
        if match:
            break
    if not match:
        return None
    customer = re.sub(r"(?i)^ST(?=\s)", "Simba", match.group(1))
    customer = re.sub(r"(?i)^Simba(?=[A-Z])", "Simba ", customer)
    return re.sub(r"\s+", " ", customer).strip()


def _release_port_country(text: str) -> tuple[str | None, str | None]:
    match = re.search(r"Port\s+of\s+discharge\s*:\s*([A-Za-z][A-Za-z .'-]{1,80})", text, re.I)
    if not match:
        return None, None
    port = re.split(r"\*|\n", match.group(1), maxsplit=1)[0].strip(" .,-").upper()
    return port or None, PORT_COUNTRY_MAP.get(port)


def _release_packing(block: str) -> tuple[float | None, float | None]:
    match = re.search(
        r"Packing\s*:\s*([\d,]+|[Oo])\s*PC\s*(?:/|\||1)\s*([\d,]+|[Oo])\s*PC",
        block,
        re.I,
    )
    if not match:
        return None, None
    return _release_num(match.group(1)), _release_num(match.group(2))


def _release_dimensions(block: str) -> tuple[float | None, float | None, float | None]:
    match = re.search(
        r"Carton\s*dimension\s*:\s*([\d.]+)\s*[x×]\s*([\d.]+)\s*[x×]\s*([\d.]+)\s*CM",
        block,
        re.I,
    )
    if not match:
        return None, None, None
    return tuple(_release_num(value) or None for value in match.groups())  # type: ignore[return-value]


def _release_volume_m3(block: str) -> float | None:
    match = re.search(r"Volume\s*:\s*[\d.]+\s*FT3\s*-\s*([\d.]+)\s*M3", block, re.I)
    return (_release_num(match.group(1)) or None) if match else None


def _inspection_date(ship_date: str | None) -> str | None:
    if not ship_date:
        return None
    try:
        value = datetime.strptime(ship_date, "%Y-%m-%d").date() - timedelta(days=3)
    except ValueError:
        return None
    while value.weekday() >= 5:
        value -= timedelta(days=1)
    return value.isoformat()


def _remarks_without_identity(remarks: str | None) -> str | None:
    if not remarks:
        return None
    value = re.sub(r"^\s*Customer\s*:\s*[A-Za-z0-9,&.' /()_-]+", "", remarks, flags=re.I)
    value = re.sub(
        r"\s*\**\s*(?:This\s+)?orders?\s+(?:is\s+)?for\s+(?:cust(?:omer)?\s*)?"
        r"(?:Simba|ST)\s*(?:Toys\s*)?"
        r"(?:Germany|Italy|Turkey|Hungary|Poland|France|Australia|UK)\s*\**[.,]?",
        "",
        value,
        flags=re.I,
    )
    value = re.sub(r"\s+", " ", value).strip(" *:-")
    return value[:4000] or None


def _customer_label_from_remarks(remarks: str | None) -> str | None:
    if remarks and re.search(r"label|barcode\s*labels?|招紙|招纸|貼紙|贴纸|利寶|利宝", remarks, re.I):
        return "欠客利宝"
    return "NO"


def _certificate_from_remarks(remarks: str | None) -> str:
    if remarks and re.search(r"certificate|證書|证书", remarks, re.I):
        return "YES"
    return "NO"


def parse_release_order_text(text: str, filename: str | None = None) -> list[dict[str, Any]] | None:
    """解析 Simba Dickie 'Release Order'（扫描件转换成 Excel/文字后）。

    转换件版式不稳定：数据可能单行、拆两行、甚至整表按"列"打散
    （所有数量一段、合同号一段、单价一段）。因此按全文扫描各类字段，
    再按出现顺序配对成行。Simba 规则：5 开头是 Master 合同、3 开头是 PO 合同。
    """
    compact = text.replace(" ", "")
    if "Mat.No" not in compact and "ReleaseOrder" not in compact:
        return None
    contract, contract_line = _release_reference(text)
    creation = re.search(r"Date\s+of\s+creation:\s*(\d{1,2}\.[A-Za-z]{3,4}\.\d{4})", text, re.I)
    po_creation_date = _release_date(creation.group(1)) if creation else None
    contact = _release_contact(text)
    remarks = _release_remarks(text)
    customer = _release_customer(remarks, text)
    port, ship_country = _release_port_country(text)

    # 货号、英文品名、包装及箱规（按 Mat.No 分块提取）。
    mats: list[dict[str, Any]] = []
    for block in re.split(r"Mat\.\s*No\.?\s*:", text)[1:]:
        mat_match = re.match(r"\s*([A-Z0-9]{6,18})", block, re.I)
        if not mat_match:
            continue
        desc_parts: list[str] = []
        for ln in [x.strip() for x in block.splitlines()][1:8]:
            if re.match(r"(Packing|Quantity)", ln, re.I):
                break
            if not ln or re.fullmatch(r"\d{12,14}", ln) or ln.startswith("Mat.EAN") or not re.search(r"[A-Za-z]{3}", ln):
                continue
            desc_parts.append(ln)
            if len(desc_parts) >= 4:
                break
        inner, outer = _release_packing(block)
        length, width, height = _release_dimensions(block)
        mats.append(
            {
                "item_no": mat_match.group(1),
                "english_name": " ".join(desc_parts)[:300],
                "inner_pack": inner,
                "outer_pack": outer,
                "outer_length_cm": length,
                "outer_width_cm": width,
                "outer_height_cm": height,
                "cbf": _release_volume_m3(block),
            }
        )

    # 只在 Quantity 表头之后扫描数据，避开表头区的下单日期等干扰。
    anchor = text.find("Quantity")
    scan = text[anchor:] if anchor >= 0 else text

    qtys: list[float] = []
    for m in re.finditer(r"([\d.,]+)\s*PC\b", scan):
        before = scan[:m.start()].rstrip()
        after = scan[m.end():].lstrip()
        if before.endswith("/") or after.startswith("/"):
            continue  # 排除 "Packing: 30 PC / 120 PC"
        qtys.append(_release_num(m.group(1)))
    masters: list[str] = []
    pos_list: list[str] = []
    for m in re.finditer(r"\b([35]\d{8})\s*/\s*(\d{1,3})\b", scan):
        (masters if m.group(1).startswith("5") else pos_list).append(f"{m.group(1)}/{m.group(2)}")
    prices: list[tuple[float, str]] = []
    service_cost = False
    for m in re.finditer(r"([\d.,]+)\s*[·\s]\s*(HKD|USD|EUR)\b", scan):
        if re.search(r"service\s*costs?", scan[max(0, m.start() - 30):m.start()], re.I):
            service_cost = True
            continue
        prices.append((_release_num(m.group(1)), m.group(2).upper()))
    dates: list[str | None] = []
    for m in re.finditer(r"(\d{1,2})[.,]\s*([A-Za-z]{3,4})[.,]\s*(\d{4})", scan):
        try:
            dates.append(datetime.strptime(f"{int(m.group(1))}.{m.group(2)[:3]}.{m.group(3)}", "%d.%b.%Y").date().isoformat())
        except ValueError:
            dates.append(None)

    if not qtys or not prices:
        return None
    counts = {len(qtys), len(prices)} | ({len(pos_list)} if pos_list else set()) | ({len(dates)} if dates else set())
    mismatch = len(counts) > 1

    # 自洽校验：总箱数 × 外箱装箱数 = 数量（Total CTN 是单据级合计，按行分箱累加比对）。
    ctn_note = None
    ctn_match = re.search(r"Total\s*CTN\s*:\s*([\d,]+)", text, re.I)
    if ctn_match and len(qtys) <= len(mats):
        outers = [mats[i]["outer_pack"] for i in range(len(qtys))]
        if all(o is not None and o > 0 for o in outers):
            expect_ctn = sum(q / o for q, o in zip(qtys, outers))
            doc_ctn = _release_num(ctn_match.group(1))
            if abs(expect_ctn - doc_ctn) > 0.51:
                ctn_note = (
                    f"箱数校验不符：数量÷装箱数合计≈{expect_ctn:.1f}箱，单据 Total CTN={doc_ctn:.0f}箱，"
                    "数量或装箱数可能识别错误，请人工核对"
                )
    rows: list[dict[str, Any]] = []
    for i, qty in enumerate(qtys):
        price, currency = prices[i] if i < len(prices) else (None, None)
        note_parts = []
        if mismatch:
            note_parts.append("转换件字段数不一致，本行为按顺序配对结果，请人工核对")
        if service_cost:
            note_parts.append("原单含 Service costs 附加费，未计入单价")
        if ctn_note:
            note_parts.append(ctn_note)
        if not customer:
            note_parts.append("PO未明确写出客名，系统未猜测，请人工补客名")
        if not contact:
            note_parts.append("PO未识别到合同联系人，请人工补充")
        if not port:
            note_parts.append("PO未识别到卸货港，走货国家留空")
        source_remark = _remarks_without_identity(remarks)
        material = mats[i] if i < len(mats) else (mats[0] if len(mats) == 1 else {})
        outer_pack = material.get("outer_pack") or None
        cartons = round(qty / outer_pack, 2) if outer_pack else None
        if isinstance(cartons, float) and cartons.is_integer():
            cartons = int(cartons)
        master_contract = masters[i] if i < len(masters) else None
        material_order_no = None
        if master_contract:
            qty_text = str(int(qty)) if float(qty).is_integer() else str(qty)
            material_order_no = f"{master_contract.split('/')[0]}-{qty_text}"
        row: dict[str, Any] = {
            "item_no": material.get("item_no"),
            # 蓝色批注规定中文产品名称从排期按货号继承，PO DESCRIPTION 写入英文名称。
            "product_name": None,
            "english_name": material.get("english_name"),
            "quantity": qty,
            "inner_pack": material.get("inner_pack"),
            "outer_pack": outer_pack,
            "cartons": cartons,
            "contract_no": contract,
            "contract_line": contract_line,
            "master_contract_no": master_contract,
            "po_contract_no": pos_list[i] if i < len(pos_list) else None,
            # 来单日期是“收到正式 PO 邮件的时间”，不能用 Date of creation 冒充。
            "order_date": None,
            "po_creation_date": po_creation_date,
            "contact": contact,
            "customer_po": None,
            "customer": customer,
            "po_ship_date": dates[i] if i < len(dates) else None,
            "inspection_date": _inspection_date(dates[i] if i < len(dates) else None),
            "shipping_mark": (
                "Simba Toys"
                if re.search(r"see\s*packaging", text, re.I)
                else (customer or "待确认（缺客名）")
            ),
            "customer_label": _customer_label_from_remarks(remarks),
            "brand": "SIMBA",
            "certificate": _certificate_from_remarks(remarks),
            "outer_length_cm": material.get("outer_length_cm"),
            "outer_width_cm": material.get("outer_width_cm"),
            "outer_height_cm": material.get("outer_height_cm"),
            "cbf": material.get("cbf"),
            "total_cbf": round(float(cartons) * float(material["cbf"]), 6)
            if cartons and material.get("cbf")
            else None,
            "ship_country": ship_country,
            "shipping_info": port,
            "material_order_no": material_order_no,
            "barcode_item_no": material.get("item_no"),
            "source_sheet": "转换PO",
            "special_remark": "；".join(filter(None, [source_remark, *note_parts])) or None,
        }
        if price is not None:
            if currency == "USD":
                row["unit_price_usd"] = price
                row["total_usd"] = round(qty * price, 2)
            else:
                row["unit_price_hkd"] = price
                row["total_hkd"] = round(qty * price, 2)
                row["currency"] = currency
        try:
            add_derived_fields(row)
        except Exception:
            pass
        rows.append(row)
    return rows or None


def collect_order_notes(text: str) -> str | None:
    notes = []
    for line in text.splitlines():
        clean = re.sub(r"\s+", " ", line).strip()
        if re.search(r"\b(Label|Shipping Mark|ORDER NOTES|REMARKS|RFID|Importer label)\b", clean, re.I):
            notes.append(clean)
    return "\n".join(notes) if notes else None


def parse_pdf_po(source: str | Path | bytes | BinaryIO, filename: str | None = None) -> dict[str, Any]:
    text, pages = read_pdf_text(source)
    warnings: list[str] = []
    if not text:
        return {
            "filename": filename or "po.pdf",
            "type": "pdf",
            "meta": {"pages": pages, "text_chars": 0},
            "rows": [],
            "warnings": ["PDF未提取到文字，可能是扫描件或图片版PDF。请上传文字版PDF、Excel PO，或在系统里手工补录。"],
            "text_preview": "",
        }

    po_no = search(r"Purchase\s+Order\s+No:\s*([A-Za-z0-9./_-]+)", text)
    order_date = parse_date_text(search(r"Order\s+Date\s+([0-9./-]+)", text))
    shipment_date = parse_date_text(
        search(r"Shipment\s+Date\s+([0-9./-]+)", text)
        or search(r"Delivery\s+Date\s+([0-9./-]+)", text)
        or search(r"Request\s+Date\s+([0-9./-]+)", text)
    )
    customer, customer_po, trade_term = parse_reference_line(text)
    remarks = collect_order_notes(text)
    rows = parse_line_items(text)
    if rows:
        # 自洽校验：逐行 数量×单价=行金额；行合计=单据TOTAL
        for row in rows:
            qty, price, total = row.get("quantity") or 0, row.get("unit_price_usd") or 0, row.get("total_usd") or 0
            if qty > 0 and price > 0 and abs(round(qty * price, 2) - total) > 0.011:
                warnings.append(
                    f"{row.get('item_no')}: 数量×单价({qty}×{price})与行金额({total})不符，请人工核对"
                )
        doc_total_m = re.search(r"TOTAL\s+NET\s+AMOUNT\s+USD?\s*([\d,]+\.\d{2})", text, re.I)
        if doc_total_m:
            line_sum = round(sum(r.get("total_usd") or 0 for r in rows), 2)
            doc_total = _release_num(doc_total_m.group(1))
            if abs(doc_total - line_sum) > 0.011:
                warnings.append(
                    f"行合计({line_sum})与单据TOTAL({doc_total})不符——可能有明细行未识别，请人工核对"
                )
    if not rows:
        release_rows = parse_release_order_text(text, filename)
        if release_rows:
            rows = release_rows
            warnings.append("已按 Simba Release Order 版式解析，请人工核对关键数值。")
    if not rows:
        warnings.append("已读取PDF文字，但未识别到标准PO明细行；请检查PDF格式或改用Excel。")
    for row in rows:
        fallback = {
            "contract_no": po_no,
            "customer": customer,
            "customer_po": customer_po,
            "order_date": order_date,
            "po_ship_date": shipment_date,
            "shipping_info": trade_term,
        }
        for field, value in fallback.items():
            if row.get(field) in (None, "") and value not in (None, ""):
                row[field] = value
        if remarks and not row.get("special_remark"):
            row["special_remark"] = remarks
        row["source_sheet"] = "PDF"
        add_derived_fields(row)
    return {
        "filename": filename or "po.pdf",
        "type": "pdf",
        "meta": {
            "pages": pages,
            "text_chars": len(text),
            "po_no": po_no,
            "customer": customer,
            "customer_po": customer_po,
            "shipment_date": shipment_date,
        },
        "rows": rows,
        "warnings": warnings,
        "text_preview": text[:1600],
    }


def parse_excel_po(source: str | Path | bytes | BinaryIO, filename: str | None = None) -> dict[str, Any]:
    if isinstance(source, (str, Path)):
        source = Path(source).read_bytes()
    try:
        result = read_schedule(source, filename=filename, allow_po_note=True)
        rows = result["records"]
        meta, summary = result["meta"], result["summary"]
        schedule_error = None
    except ValueError as exc:
        rows, meta, summary, schedule_error = [], {}, {}, exc

    warnings: list[str] = []
    if not rows:
        # 不是排期表版式时，按"扫描件转换成 Excel 的 Release Order"兜底解析。
        release_rows = parse_release_order_text(workbook_text(source), filename)
        if release_rows:
            rows = release_rows
            warnings.append("已按 Simba Release Order 转换版式解析；转换件可能有识别误差，关键数量/合同号请人工核对。")
        elif schedule_error is not None:
            raise schedule_error
        else:
            warnings.append("Excel中未找到可识别的PO明细行。")
    return {
        "filename": filename or "po.xlsx",
        "type": "excel",
        "meta": meta,
        "summary": summary,
        "rows": rows,
        "warnings": warnings,
        "text_preview": "",
    }


def parse_po_file(source: str | Path | bytes | BinaryIO, filename: str | None = None) -> dict[str, Any]:
    name = filename or (Path(source).name if isinstance(source, (str, Path)) else "upload")
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return parse_pdf_po(source, filename=name)
    if suffix in {".xlsx", ".xlsm"}:
        return parse_excel_po(source, filename=name)
    raise ValueError(f"Unsupported PO file type: {suffix or 'unknown'}")
