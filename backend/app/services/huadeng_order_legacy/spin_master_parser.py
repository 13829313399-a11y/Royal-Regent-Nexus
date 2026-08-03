from __future__ import annotations

import csv
import io
import json
import os
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pdfplumber
import xlrd
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


BASE_DIR = Path(__file__).resolve().parent
LEGACY_DATA_DIR = BASE_DIR / "华登Spin Master客AI"
LEGACY_EXCEL_NAME = "2026年Spin Master -AI.xls"
LEGACY_PDF_NAME = "PRD-ROYAL REGENT PRODUCT-4500706644.pdf"
UPLOAD_STEM = "current_schedule"
UPLOAD_METADATA_NAME = "current_schedule.json"

# Kept as compatibility aliases for existing scripts and local data.
DATA_DIR = LEGACY_DATA_DIR
EXCEL_PATH = LEGACY_DATA_DIR / LEGACY_EXCEL_NAME
PDF_PATH = LEGACY_DATA_DIR / LEGACY_PDF_NAME

SCHEDULE_SHEET = "SPIN排期"
SUMMARY_SHEET = "SPIN总汇"


@dataclass(frozen=True)
class ColumnDef:
    index: int
    key: str
    title: str
    kind: str = "text"


COLUMNS = [
    ColumnDef(1, "order_date", "来单期", "date"),
    ColumnDef(2, "change_note", "改单期"),
    ColumnDef(3, "owner", "联系人"),
    ColumnDef(4, "customer_po", "客PO", "id"),
    ColumnDef(5, "contract_no", "合同号", "id"),
    ColumnDef(6, "customer", "客名"),
    ColumnDef(7, "version", "版本"),
    ColumnDef(8, "item_si", "货号/SI", "id"),
    ColumnDef(9, "product_name", "产品名称"),
    ColumnDef(10, "quantity", "数量", "number"),
    ColumnDef(11, "inner_carton", "内箱", "number"),
    ColumnDef(12, "outer_carton", "外箱", "number"),
    ColumnDef(13, "total_cartons", "总箱数", "number"),
    ColumnDef(14, "special_note", "特别备注"),
    ColumnDef(15, "box_mark", "箱唛"),
    ColumnDef(16, "customer_supplied_label", "客供利宝"),
    ColumnDef(17, "customer_supplied_material", "客供物料"),
    ColumnDef(18, "material_reply_date", "来料复期", "date"),
    ColumnDef(19, "plastic_reply_date", "胶件复期", "date"),
    ColumnDef(20, "carton_reply_date", "纸箱复期", "date"),
    ColumnDef(21, "injection_reply_date", "注料复期", "date"),
    ColumnDef(22, "pull_date", "上拉期", "date"),
    ColumnDef(23, "completion_date", "完成期", "date"),
    ColumnDef(24, "date_code", "日期码", "id"),
    ColumnDef(25, "workshop_recheck", "车间复验"),
    ColumnDef(26, "inspection_date", "验货期", "date"),
    ColumnDef(27, "po_ship_date", "PO走货期", "date"),
    ColumnDef(28, "booking_lead_days", "预计订仓时间", "number"),
    ColumnDef(29, "receiving_port", "收货地点/装箱港"),
    ColumnDef(30, "unit_hkd", "单价港币", "number"),
    ColumnDef(31, "unit_usd", "单价(美金)", "number"),
    ColumnDef(32, "total_hkd", "总货价(港币)", "number"),
    ColumnDef(33, "total_usd", "总货价(美金)", "number"),
    ColumnDef(34, "assembly_cost_hkd", "报价装配工HK$", "number"),
    ColumnDef(35, "english_name", "英文品名"),
    ColumnDef(36, "material_follow_up_date", "预追客供物料期", "date"),
    ColumnDef(37, "certificate", "证书"),
    ColumnDef(38, "outer_length", "外箱 长", "number"),
    ColumnDef(39, "outer_width", "外箱 宽", "number"),
    ColumnDef(40, "outer_height", "外箱 高", "number"),
    ColumnDef(41, "unit_cbm", "单箱CBM", "number"),
    ColumnDef(42, "total_cbm", "总CBM", "number"),
    ColumnDef(43, "gross_weight_kg", "毛重KG", "number"),
    ColumnDef(44, "net_weight_kg", "净重KG", "number"),
    ColumnDef(45, "ship_country", "走货国家"),
    ColumnDef(46, "so_no", "SO.NO", "id"),
    ColumnDef(47, "warehouse_delivery_date", "交仓日期", "date"),
    ColumnDef(48, "truck_detail", "车明细"),
    ColumnDef(49, "brand", "品牌"),
    ColumnDef(50, "entered_system", "是否已入系统", "date"),
    ColumnDef(51, "inspection_ok", "验货是否OK"),
    ColumnDef(52, "shipping_notice_time", "通知出货时间", "date"),
    ColumnDef(53, "closing_time", "结关时间", "date"),
    ColumnDef(54, "prep_order_no", "备料单号", "id"),
    ColumnDef(55, "invoice_no", "发票号码", "id"),
    ColumnDef(56, "invoice_date", "发票日期", "date"),
    ColumnDef(57, "weilian", "伟联"),
    ColumnDef(61, "tattoo_supplier", "纹身贴纸供应商"),
    ColumnDef(62, "transfer_film_supplier", "转移膜供应商"),
    ColumnDef(63, "sand", "沙"),
    ColumnDef(64, "imported_ic", "进口IC"),
    ColumnDef(65, "rfid_fee", "RFID利宝收费"),
    ColumnDef(66, "amazon_bag_fee", "AMAZON胶袋收费"),
    ColumnDef(67, "inspection_report_no", "验货报告编号", "id"),
    ColumnDef(68, "date_code_result", "日期码", "id"),
    ColumnDef(69, "inspection_result", "验货结果"),
]

DATE_COLUMNS = {col.index for col in COLUMNS if col.kind == "date"}
TEXT_ID_COLUMNS = {col.index for col in COLUMNS if col.kind == "id"}
COLUMN_BY_KEY = {col.key: col for col in COLUMNS}


def _configured_path(name: str, default: Path) -> Path:
    raw = os.environ.get(name, "").strip()
    path = Path(raw).expanduser() if raw else default
    if not path.is_absolute():
        path = BASE_DIR / path
    return path.resolve()


def data_dir() -> Path:
    return _configured_path("SPIN_DATA_DIR", LEGACY_DATA_DIR)


def upload_dir() -> Path:
    return _configured_path("SPIN_UPLOAD_DIR", BASE_DIR / "uploads")


def active_excel_path() -> Path:
    explicit = os.environ.get("SPIN_EXCEL_PATH", "").strip()
    if explicit:
        return _configured_path("SPIN_EXCEL_PATH", EXCEL_PATH)

    candidates = [upload_dir() / f"{UPLOAD_STEM}.xls", upload_dir() / f"{UPLOAD_STEM}.xlsx"]
    existing = [path for path in candidates if path.exists()]
    if existing:
        return max(existing, key=lambda path: path.stat().st_mtime_ns)
    return data_dir() / LEGACY_EXCEL_NAME


def active_pdf_path() -> Path:
    return _configured_path("SPIN_PDF_PATH", data_dir() / LEGACY_PDF_NAME)


def schedule_file_info(excel_path: Path | None = None) -> dict[str, Any]:
    path = Path(excel_path or active_excel_path())
    info: dict[str, Any] = {
        "path": str(path),
        "name": path.name,
        "original_name": path.name,
        "exists": path.exists(),
        "size": path.stat().st_size if path.exists() else 0,
        "modified_at": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds") if path.exists() else "",
        "uploaded_at": "",
    }
    metadata_path = upload_dir() / UPLOAD_METADATA_NAME
    if metadata_path.exists():
        try:
            metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
            if metadata.get("stored_name") == path.name:
                info.update(
                    {
                        "original_name": metadata.get("original_name") or path.name,
                        "uploaded_at": metadata.get("uploaded_at") or "",
                    }
                )
        except (OSError, ValueError, TypeError):
            pass
    return info


def workbook_paths() -> dict[str, Any]:
    excel_path = active_excel_path()
    pdf_path = active_pdf_path()
    return {
        "data_dir": str(data_dir()),
        "upload_dir": str(upload_dir()),
        "excel": str(excel_path),
        "pdf": str(pdf_path),
        "source": schedule_file_info(excel_path),
    }


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"[ \t]+", " ", str(value).replace("\r", "\n")).strip()


def _parse_xlrd_date(value: Any, datemode: int) -> str:
    if not isinstance(value, (int, float)) or value <= 0:
        return ""
    try:
        return xlrd.xldate_as_datetime(value, datemode).date().isoformat()
    except (OverflowError, ValueError, TypeError):
        return ""


def _number(value: Any) -> float | int | None:
    if value in ("", None):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return int(value)
        return round(float(value), 6)
    text = _clean_text(value).replace(",", "")
    try:
        number = float(text)
    except ValueError:
        return None
    if number.is_integer():
        return int(number)
    return round(number, 6)


def _cell_value(value: Any, col: ColumnDef, datemode: int | None = None) -> Any:
    if value in ("", None):
        return ""
    if col.index in DATE_COLUMNS:
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if datemode is not None and isinstance(value, (int, float)) and value > 20000:
            return _parse_xlrd_date(value, datemode)
        if datemode is None and isinstance(value, (int, float)) and value > 20000:
            try:
                return from_excel(value).date().isoformat()
            except (OverflowError, ValueError, TypeError):
                return ""
        return _clean_text(value)
    if col.kind == "number":
        number = _number(value)
        return "" if number is None else number
    if col.index in TEXT_ID_COLUMNS:
        return _clean_text(value)
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, float):
        return round(value, 6)
    return _clean_text(value)


def _contract_parts(contract_no: str) -> tuple[str, int | None]:
    match = re.search(r"(\d+)\s*/\s*(\d+)", contract_no or "")
    if not match:
        return contract_no or "", None
    return match.group(1), int(match.group(2))


def _item_parts(item_si: str) -> dict[str, str]:
    groups = re.findall(r"\d+", item_si or "")
    parts = {
        "material_group": "",
        "material_no": "",
        "sales_material": "",
        "sales_order": "",
        "sales_line_item": "",
    }
    if len(groups) >= 1:
        parts["material_group"] = groups[0]
    if len(groups) >= 2:
        parts["material_no"] = groups[1]
    if len(groups) >= 3:
        parts["sales_material"] = groups[2]
    if len(groups) >= 5:
        parts["sales_order"] = groups[-2]
        parts["sales_line_item"] = groups[-1]
    return parts


def _row_status(row: dict[str, Any], today: date | None = None) -> str:
    today = today or date.today()
    material_text = " ".join(
        str(row.get(key, ""))
        for key in ("box_mark", "customer_supplied_label", "customer_supplied_material")
    )
    if "欠" in material_text:
        return "缺资料"

    ship_date = _to_date(row.get("po_ship_date"))
    inspection_date = _to_date(row.get("inspection_date"))
    if inspection_date and ship_date and inspection_date > ship_date:
        return "交期风险"
    if ship_date:
        days = (ship_date - today).days
        if days < 0:
            return "已过走货期"
        if days <= 14:
            return "临近走货"
    return "正常"


def _to_date(value: Any) -> date | None:
    if not value:
        return None
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").date()
    except ValueError:
        return None


def _schedule_row(
    values: list[Any] | tuple[Any, ...],
    excel_row: int,
    datemode: int | None,
    column_offset: int = 1,
) -> dict[str, Any] | None:
    row = {
        col.key: _cell_value(
            values[col.index - 1 + column_offset]
            if col.index - 1 + column_offset < len(values)
            else "",
            col,
            datemode,
        )
        for col in COLUMNS
    }
    if not _is_schedule_row(row):
        return None
    contract_base, contract_line = _contract_parts(str(row.get("contract_no", "")))
    row.update(
        {
            "excel_row": excel_row,
            "row_id": f"{SCHEDULE_SHEET}-{excel_row}",
            "contract_base": contract_base,
            "contract_line": contract_line,
            "status": _row_status(row),
        }
    )
    row.update(_item_parts(str(row.get("item_si", ""))))
    return row


def load_schedule(sheet_name: str = SCHEDULE_SHEET, excel_path: Path | str | None = None) -> dict[str, Any]:
    source = Path(excel_path or active_excel_path())
    if not source.exists():
        raise FileNotFoundError(f"Excel file not found: {source}")
    if source.suffix.lower() not in {".xls", ".xlsx"}:
        raise ValueError("排期文件仅支持 .xls 或 .xlsx")

    headers = [
        {"index": col.index, "key": col.key, "title": col.title, "kind": col.kind}
        for col in COLUMNS
    ]
    rows: list[dict[str, Any]] = []

    if source.suffix.lower() == ".xls":
        workbook = xlrd.open_workbook(str(source))
        if sheet_name not in workbook.sheet_names():
            sheet_name = next(
                (
                    name for name in workbook.sheet_names()
                    if max(
                        sum(
                            _clean_text(workbook.sheet_by_name(name).cell_value(0, c + offset)) == col.title
                            for c, col in enumerate(COLUMNS[:15])
                            if c + offset < workbook.sheet_by_name(name).ncols
                        )
                        for offset in (0, 1)
                    ) >= 6
                ),
                sheet_name,
            )
        if sheet_name not in workbook.sheet_names():
            raise ValueError(f"Excel 中缺少工作表：{SCHEDULE_SHEET}")
        sheet = workbook.sheet_by_name(sheet_name)
        offset = 0 if _clean_text(sheet.cell_value(0, 0)) == COLUMNS[0].title else 1
        for row_index in range(1, sheet.nrows):
            row = _schedule_row(sheet.row_values(row_index), row_index + 1, workbook.datemode, offset)
            if row:
                rows.append(row)
    else:
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                sheet_name = next(
                    (
                        name for name in workbook.sheetnames
                        if max(
                            sum(
                                _clean_text(workbook[name].cell(1, c + 1 + offset).value) == col.title
                                for c, col in enumerate(COLUMNS[:15])
                            )
                            for offset in (0, 1)
                        ) >= 6
                    ),
                    sheet_name,
                )
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"Excel 中缺少工作表：{SCHEDULE_SHEET}")
            sheet = workbook[sheet_name]
            offset = 0 if _clean_text(sheet.cell(1, 1).value) == COLUMNS[0].title else 1
            for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                row = _schedule_row(values, excel_row, None, offset)
                if row:
                    rows.append(row)
        finally:
            workbook.close()

    return {
        "sheet": sheet_name,
        "source": str(source),
        "headers": headers,
        "rows": rows,
        "stats": build_stats(rows),
    }


def validate_schedule_file(excel_path: Path | str) -> dict[str, Any]:
    schedule = load_schedule(excel_path=excel_path)
    if not schedule["rows"]:
        raise ValueError("没有从 SPIN排期 工作表读取到有效排期记录")
    return schedule


def _is_schedule_row(row: dict[str, Any]) -> bool:
    identity = [
        row.get("contract_no"),
        row.get("customer_po"),
        row.get("customer"),
        row.get("item_si"),
        row.get("product_name"),
    ]
    if not any(identity):
        return False
    if row.get("quantity") in ("", None) and not row.get("contract_no"):
        return False
    if row.get("contract_no") in (".", "合计", "小计"):
        return False
    return True


def build_stats(rows: list[dict[str, Any]]) -> dict[str, Any]:
    total_quantity = sum(float(row.get("quantity") or 0) for row in rows)
    total_cartons = sum(float(row.get("total_cartons") or 0) for row in rows)
    total_usd = sum(float(row.get("total_usd") or 0) for row in rows)
    customers = sorted({str(row.get("customer", "")).replace("\n", " ") for row in rows if row.get("customer")})
    statuses: dict[str, int] = {}
    ship_months: dict[str, float] = {}
    customer_quantity: dict[str, float] = {}

    for row in rows:
        statuses[row["status"]] = statuses.get(row["status"], 0) + 1
        ship_date = str(row.get("po_ship_date") or "")
        month = ship_date[:7] if len(ship_date) >= 7 else "未定"
        ship_months[month] = ship_months.get(month, 0) + float(row.get("quantity") or 0)
        customer = str(row.get("customer") or "未填").replace("\n", " ")
        customer_quantity[customer] = customer_quantity.get(customer, 0) + float(row.get("quantity") or 0)

    top_customers = sorted(
        [{"name": name, "quantity": round(quantity, 2)} for name, quantity in customer_quantity.items()],
        key=lambda item: item["quantity"],
        reverse=True,
    )[:10]

    return {
        "row_count": len(rows),
        "customer_count": len(customers),
        "total_quantity": round(total_quantity, 2),
        "total_cartons": round(total_cartons, 2),
        "total_usd": round(total_usd, 2),
        "statuses": statuses,
        "ship_months": dict(sorted(ship_months.items())),
        "top_customers": top_customers,
        "customers": customers,
    }


def run_quality_checks(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    checks = [
        _check_required_fields(rows),
        _check_carton_math(rows),
        _check_duplicate_contracts(rows),
        _check_date_sequence(rows),
        _check_material_status(rows),
    ]
    return checks


def _issue(row: dict[str, Any], message: str, field: str = "") -> dict[str, Any]:
    return {
        "excel_row": row.get("excel_row"),
        "contract_no": row.get("contract_no"),
        "customer_po": row.get("customer_po"),
        "field": field,
        "message": message,
    }


def _check_required_fields(rows: list[dict[str, Any]]) -> dict[str, Any]:
    required = [
        ("order_date", "来单期"),
        ("owner", "联系人"),
        ("contract_no", "合同号"),
        ("customer", "客名"),
        ("item_si", "货号/SI"),
        ("product_name", "产品名称"),
        ("quantity", "数量"),
        ("outer_carton", "外箱"),
        ("total_cartons", "总箱数"),
        ("po_ship_date", "PO走货期"),
    ]
    issues = []
    for row in rows:
        missing = [label for key, label in required if row.get(key) in ("", None)]
        if missing:
            issues.append(_issue(row, "缺少：" + "、".join(missing), "required"))
    return _quality_result("必填字段", issues, "检查核心排期字段是否完整")


def _check_carton_math(rows: list[dict[str, Any]]) -> dict[str, Any]:
    issues = []
    for row in rows:
        quantity = _number(row.get("quantity"))
        outer = _number(row.get("outer_carton"))
        total = _number(row.get("total_cartons"))
        if not quantity or not outer or total in (None, ""):
            continue
        expected = float(quantity) / float(outer)
        if abs(expected - float(total)) > 0.05:
            issues.append(
                _issue(
                    row,
                    f"数量/外箱={expected:.2f}，表内总箱数={float(total):.2f}",
                    "total_cartons",
                )
            )
    return _quality_result("箱数计算", issues, "检查总箱数是否等于数量除以外箱")


def _check_duplicate_contracts(rows: list[dict[str, Any]]) -> dict[str, Any]:
    seen: dict[str, dict[str, Any]] = {}
    issues = []
    for row in rows:
        contract = str(row.get("contract_no") or "")
        if not contract:
            continue
        if contract in seen:
            issues.append(_issue(row, f"与第 {seen[contract].get('excel_row')} 行合同号重复", "contract_no"))
        else:
            seen[contract] = row
    return _quality_result("合同重复", issues, "检查合同号/行号是否重复")


def _check_date_sequence(rows: list[dict[str, Any]]) -> dict[str, Any]:
    issues = []
    for row in rows:
        inspection = _to_date(row.get("inspection_date"))
        ship = _to_date(row.get("po_ship_date"))
        if inspection and ship and inspection > ship:
            issues.append(_issue(row, "验货期晚于 PO 走货期", "inspection_date"))
    return _quality_result("日期顺序", issues, "检查验货期是否早于走货期")


def _check_material_status(rows: list[dict[str, Any]]) -> dict[str, Any]:
    issues = []
    for row in rows:
        text = " ".join(
            str(row.get(key, ""))
            for key in ("box_mark", "customer_supplied_label", "customer_supplied_material")
        )
        if "欠" in text:
            issues.append(_issue(row, "箱唛/利宝/客供物料仍有欠缺", "material"))
    return _quality_result("资料状态", issues, "检查箱唛、利宝、客供物料欠缺项")


def _quality_result(name: str, issues: list[dict[str, Any]], description: str) -> dict[str, Any]:
    return {
        "name": name,
        "description": description,
        "status": "pass" if not issues else "warn",
        "issue_count": len(issues),
        "issues": issues[:50],
    }


def parse_po_pdf(pdf_path: Path = PDF_PATH) -> dict[str, Any]:
    if not pdf_path.exists():
        raise FileNotFoundError(f"PDF file not found: {pdf_path}")

    pages: list[str] = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    text = "\n".join(pages)

    po_number = _first_match(r"PURCHASE ORDER\s+.*?NUMBER PAGE\s*\n(?P<value>\d+)", text, flags=re.S)
    if not po_number:
        po_number = _first_match(r"\b(45\d{8})\b", text)
    order_date = _format_us_date(_first_match(r"\b(\d{2}/\d{2}/\d{4})\b", text))
    vendor = _first_match(r"ORDERED FROM:\s*\d+\s+SHIP TO\s*\n(?P<value>.+?)\s+Spin Master", text, flags=re.S)
    buyer = _first_match(r"\n(?P<value>Karen Sun)\s*\n", text)

    item_matches = list(
        re.finditer(r"(?m)^(?P<line>\d{1,3})\s+(?P<qty>[\d,]+\.\d{2})\s+PC\s+(?P<material_no>\d+)\s*$", text)
    )
    items = []
    for index, match in enumerate(item_matches):
        block_end = item_matches[index + 1].start() if index + 1 < len(item_matches) else len(text)
        block = text[match.start() : block_end]
        items.append(_parse_pdf_item(block, match))

    return {
        "source": str(pdf_path),
        "po_number": po_number,
        "order_date": order_date,
        "vendor": _compact(vendor),
        "buyer": buyer,
        "page_count": len(pages),
        "item_count": len(items),
        "items": items,
        "total_quantity": round(sum(float(item.get("quantity") or 0) for item in items), 2),
        "total_net_usd": round(sum(float(item.get("net_usd") or 0) for item in items), 2),
    }


def _spreadsheet_rows(path: Path) -> tuple[list[list[Any]], int]:
    """Read WPS PDF-to-Excel output without treating it as a schedule workbook."""
    suffix = path.suffix.lower()
    if suffix == ".xls":
        book = xlrd.open_workbook(path)
        rows: list[list[Any]] = []
        for sheet in book.sheets():
            for row_index in range(sheet.nrows):
                values: list[Any] = []
                for column_index in range(sheet.ncols):
                    cell = sheet.cell(row_index, column_index)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                    values.append(value)
                rows.append(values)
        return rows, len(book.sheets())

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [
            list(row)
            for sheet in book.worksheets
            for row in sheet.iter_rows(values_only=True)
        ]
        return rows, len(book.worksheets)
    finally:
        book.close()


def _excel_cell_text(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"[ \t]+", " ", str(value).replace("\r", "\n")).strip()


def _excel_number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"[-+]?\d[\d.,]*", _excel_cell_text(value))
    if not match:
        return None
    text = match.group(0)
    if "." in text and "," in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        head, _, tail = text.rpartition(",")
        text = head.replace(",", "") + "." + tail if len(tail) == 2 and head else text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _excel_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _excel_cell_text(value)
    for pattern in ("%m/%d/%Y", "%Y-%m-%d", "%d-%b-%Y", "%d.%b.%Y"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    match = re.search(r"\b(\d{1,2}/\d{1,2}/\d{4})\b", text)
    return _format_us_date(match.group(1)) if match else ""


def _merged_excel_items(rows: list[list[Any]]) -> list[dict[str, Any]]:
    """Handle WPS conversions that merge two or more PO lines into one tall cell."""
    columns: dict[str, int] = {}
    line_numbers: list[int] = []
    quantities: list[float] = []
    detail_parts: list[str] = []
    prices: list[float] = []
    deliveries: list[str] = []
    nets: list[float] = []

    for raw_row in rows:
        values = [_excel_cell_text(value) for value in raw_row]
        is_header = any(
            "MAT. NUMBER" in value.upper() and "DESCRIPTION" in value.upper()
            for value in values
        )
        if is_header:
            for index, value in enumerate(values):
                upper = re.sub(r"\s+", " ", value.upper()).strip()
                if upper == "LINE NO":
                    columns["line"] = index
                elif "QUANTITY" in upper and "ORDERED" in upper:
                    columns["quantity"] = index
                elif upper == "U/M":
                    columns["unit"] = index
                elif "MAT. NUMBER" in upper:
                    columns["detail"] = index
                elif "PRICE/UNIT" in upper:
                    columns["price"] = index
                elif "DEL. DATE" in upper:
                    columns["delivery"] = index
                elif upper == "NET VALUE":
                    columns["net"] = index
            continue

        if not {"line", "quantity", "detail", "price", "delivery", "net"}.issubset(columns):
            continue

        line_cell = values[columns["line"]] if columns["line"] < len(values) else ""
        quantity_cell = values[columns["quantity"]] if columns["quantity"] < len(values) else ""
        detail_cell = values[columns["detail"]] if columns["detail"] < len(values) else ""
        price_cell = values[columns["price"]] if columns["price"] < len(values) else ""
        delivery_cell = values[columns["delivery"]] if columns["delivery"] < len(values) else ""
        net_cell = values[columns["net"]] if columns["net"] < len(values) else ""

        line_numbers.extend(
            int(line)
            for line in line_cell.splitlines()
            if re.fullmatch(r"\d{1,3}", line.strip()) and int(line) % 10 == 0
        )
        quantities.extend(
            number
            for line in quantity_cell.splitlines()
            if re.fullmatch(r"[\d.,]+", line.strip())
            for number in [_excel_number(line)]
            if number is not None
        )
        if detail_cell:
            detail_parts.append(detail_cell)
        for match in re.finditer(
            r"([\d.,]+)\s+USD\s+Per\s+([\d.,]+)\s*(?:PC|PCS|EA|EACH|SET)",
            price_cell,
            re.I | re.S,
        ):
            value = _excel_number(match.group(1))
            denominator = _excel_number(match.group(2))
            if value is not None and denominator not in (None, 0):
                prices.append(value * 1000 / denominator)
        deliveries.extend(
            _format_us_date(value)
            for value in re.findall(r"\b\d{1,2}/\d{1,2}/\d{4}\b", delivery_cell)
        )
        nets.extend(
            number
            for line in net_cell.splitlines()
            if re.fullmatch(r"[\d.,]+", line.strip())
            for number in [_excel_number(line)]
            if number is not None
        )

    if not line_numbers:
        return []
    detail_text = "\n".join(detail_parts)
    product_matches = list(re.finditer(r"(?m)^(\d{7,})\s*$", detail_text))
    chunks: list[tuple[str, str]] = []
    for index, match in enumerate(product_matches):
        end = product_matches[index + 1].start() if index + 1 < len(product_matches) else len(detail_text)
        chunks.append((match.group(1), detail_text[match.end():end]))
    item_count = min(len(line_numbers), len(quantities), len(chunks))
    if item_count == 0:
        return []
    line_numbers = line_numbers[:item_count]
    quantities = quantities[:item_count]

    items: list[dict[str, Any]] = []
    for index, line_no in enumerate(line_numbers):
        material_no, detail = chunks[index]
        lines = [_compact(line) for line in detail.splitlines() if _compact(line)]
        description_lines: list[str] = []
        for line in lines:
            if re.match(r"(?:Sales Material|Material Group|We provide|Sales Order|Customer|Customer PO)\s*:", line, re.I):
                break
            description_lines.append(line)
        sales_order_match = re.search(
            r"Sales Order:\s*(?P<sales_order>\d+)\s+Line Item:\s*(?P<line_item>\d+)",
            detail,
            re.I,
        )
        items.append(
            {
                "line_no": line_no,
                "quantity": quantities[index],
                "material_no": material_no,
                "description": _compact(" ".join(description_lines[:3])),
                "sales_material": _first_match(r"Sales Material:\s*(?P<value>\d+)", detail, flags=re.I),
                "price_per_1000_usd": prices[index] if index < len(prices) else None,
                "delivery_date": deliveries[index] if index < len(deliveries) else "",
                "net_usd": nets[index] if index < len(nets) else None,
                "material_group": _first_match(r"Material Group:\s*(?P<value>\d+)", detail, flags=re.I),
                "sales_order": sales_order_match.group("sales_order") if sales_order_match else "",
                "sales_line_item": sales_order_match.group("line_item") if sales_order_match else "",
                "customer": _compact(
                    _first_match(
                        r"Customer:\s*(?P<value>.*?)(?:\nCustomer PO:|Customer PO:)",
                        detail,
                        flags=re.I | re.S,
                    )
                ),
                "customer_po": _compact(
                    _first_match(r"Customer PO:\s*(?P<value>[^\r\n]+)", detail, flags=re.I)
                ),
            }
        )
    return items


def parse_po_excel(excel_path: Path) -> dict[str, Any]:
    """Parse the real Spin Master WPS conversion layout cell-by-cell."""
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    rows, sheet_count = _spreadsheet_rows(excel_path)
    row_texts = [
        [_excel_cell_text(value) for value in row if _excel_cell_text(value)]
        for row in rows
    ]
    flat_text = "\n".join("  ".join(row) for row in row_texts if row)

    po_number = ""
    order_date = ""
    vendor = ""
    buyer = ""
    max_page = 0
    for row in row_texts:
        joined = "\n".join(row)
        if not po_number:
            match = re.search(r"\b(45\d{8})\b", joined)
            if match:
                po_number = match.group(1)
        if not order_date:
            for value in row:
                parsed_date = _excel_date(value)
                if parsed_date:
                    order_date = parsed_date
                    break
        if not vendor:
            match = re.search(r"ORDERED FROM:\s*\d+\s*[\r\n| ]+([^\r\n|]+)", joined, re.I)
            if match:
                vendor = _compact(match.group(1))
        if not buyer and "Karen Sun" in joined:
            buyer = "Karen Sun"
        page_match = re.search(r"Page\s+(\d+)\s+of\s+(\d+)", joined, re.I)
        if page_match:
            max_page = max(max_page, int(page_match.group(2)))

    items: list[dict[str, Any]] = []
    for raw_row in rows:
        values = [_excel_cell_text(value) for value in raw_row]
        nonempty = [(index, value) for index, value in enumerate(values) if value]
        if len(nonempty) < 4:
            continue
        line_text = nonempty[0][1]
        quantity_text = nonempty[1][1]
        unit_text = nonempty[2][1].upper()
        if not re.fullmatch(r"\d{1,3}", line_text) or unit_text not in {"PC", "PCS", "EA", "EACH", "SET"}:
            continue
        quantity = _excel_number(quantity_text)
        if quantity is None:
            continue

        detail = nonempty[3][1]
        material_match = re.match(r"\s*([A-Za-z0-9.-]{4,})\b", detail)
        if not material_match:
            continue
        material_no = material_match.group(1)
        detail_lines = [_compact(line) for line in detail.splitlines() if _compact(line)]
        description_lines: list[str] = []
        for index, line in enumerate(detail_lines):
            if index == 0:
                remainder = _compact(re.sub(rf"^{re.escape(material_no)}\b", "", line))
                if remainder:
                    description_lines.append(remainder)
                continue
            if re.match(r"(?:Sales Material|Material Group|We provide|Sales Order|Customer|Customer PO)\s*:", line, re.I):
                break
            description_lines.append(line)
        description = _compact(" ".join(description_lines[:3]))

        sales_material = _first_match(r"Sales Material:\s*(?P<value>\d+)", detail, flags=re.I)
        material_group = _first_match(r"Material Group:\s*(?P<value>\d+)", detail, flags=re.I)
        sales_order_match = re.search(
            r"Sales Order:\s*(?P<sales_order>\d+)\s+Line Item:\s*(?P<line_item>\d+)",
            detail,
            re.I,
        )
        customer = _first_match(
            r"Customer:\s*(?P<value>.*?)(?:\nCustomer PO:|Customer PO:)",
            detail,
            flags=re.I | re.S,
        )
        customer_po = _first_match(r"Customer PO:\s*(?P<value>[^\r\n]+)", detail, flags=re.I)

        remaining = [value for _, value in nonempty[4:]]
        remainder_text = "\n".join(remaining)
        price_match = re.search(
            r"([\d.,]+)\s+USD\s+Per\s+([\d.,]+)\s*(?:PC|PCS|EA|EACH|SET)",
            remainder_text,
            re.I | re.S,
        )
        delivery_date = next((_excel_date(value) for value in remaining if _excel_date(value)), "")
        numeric_tail = [
            _excel_number(value)
            for value in remaining
            if not _excel_date(value) and not re.search(r"\bUSD\s+Per\b", value, re.I)
        ]
        numeric_tail = [value for value in numeric_tail if value is not None]
        net_usd = numeric_tail[-1] if numeric_tail else None
        price_per_qty = _excel_number(price_match.group(2)) if price_match else None
        price_value = _excel_number(price_match.group(1)) if price_match else None
        price_per_1000 = (
            price_value * 1000 / price_per_qty
            if price_value is not None and price_per_qty not in (None, 0)
            else None
        )

        items.append(
            {
                "line_no": int(line_text),
                "quantity": quantity,
                "material_no": material_no,
                "description": description,
                "sales_material": sales_material,
                "price_per_1000_usd": price_per_1000,
                "delivery_date": delivery_date,
                "net_usd": net_usd,
                "material_group": material_group,
                "sales_order": sales_order_match.group("sales_order") if sales_order_match else "",
                "sales_line_item": sales_order_match.group("line_item") if sales_order_match else "",
                "customer": _compact(customer),
                "customer_po": _compact(customer_po),
            }
        )

    structured_items = _merged_excel_items(rows)
    if structured_items and len(structured_items) >= len(items):
        items = structured_items
    if not items:
        raise ValueError("WPS 转换 Excel 中未识别到 Spin Master 产品明细表头或数据行。")
    return {
        "source": str(excel_path),
        "po_number": po_number,
        "order_date": order_date,
        "vendor": vendor,
        "buyer": buyer,
        "page_count": max_page or sheet_count,
        "item_count": len(items),
        "items": items,
        "total_quantity": round(sum(float(item.get("quantity") or 0) for item in items), 2),
        "total_net_usd": round(sum(float(item.get("net_usd") or 0) for item in items), 2),
        "source_format": excel_path.suffix.lower().lstrip("."),
        "raw_text": flat_text[:10000],
    }


def parse_po_file(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_po_pdf(path)
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        return parse_po_excel(path)
    raise ValueError(f"Spin Master PO 不支持 {suffix or '无扩展名'}")


def _parse_pdf_item(block: str, match: re.Match[str]) -> dict[str, Any]:
    sales = re.search(
        r"Sales Material:\s*(?P<sales_material>\d+)\s+"
        r"(?P<price>[\d,]+\.\d{2})\s+USD Per\s+"
        r"(?P<delivery>\d{2}/\d{2}/\d{4})\s+"
        r"(?P<net>[\d,]+\.\d{2})",
        block,
    )
    material_group = _first_match(r"Material Group:\s*(?P<value>\d+)", block)
    sales_order = re.search(r"Sales Order:\s*(?P<sales_order>\d+)\s+Line Item:\s*(?P<line_item>\d+)", block)
    customer = _first_match(r"Customer:\s*(?P<value>.*?)(?:\nCustomer PO:|Customer PO:)", block, flags=re.S)
    customer_po = _first_match(r"Customer PO:\s*(?P<value>[^\n]+)", block)
    description = _parse_description(block)

    item = {
        "line_no": int(match.group("line")),
        "quantity": float(match.group("qty").replace(",", "")),
        "material_no": match.group("material_no"),
        "description": description,
        "sales_material": "",
        "price_per_1000_usd": None,
        "delivery_date": "",
        "net_usd": None,
        "material_group": material_group,
        "sales_order": "",
        "sales_line_item": "",
        "customer": _compact(customer),
        "customer_po": _compact(customer_po),
    }
    if sales:
        item.update(
            {
                "sales_material": sales.group("sales_material"),
                "price_per_1000_usd": float(sales.group("price").replace(",", "")),
                "delivery_date": _format_us_date(sales.group("delivery")),
                "net_usd": float(sales.group("net").replace(",", "")),
            }
        )
    if sales_order:
        item["sales_order"] = sales_order.group("sales_order")
        item["sales_line_item"] = sales_order.group("line_item")
    return item


def _parse_description(block: str) -> str:
    lines = [line.strip() for line in block.splitlines()]
    if len(lines) < 2:
        return ""
    description_lines = []
    for line in lines[1:]:
        if line.startswith("Sales Material:"):
            break
        if not line or line.startswith("PURCHASE ORDER"):
            continue
        if line.startswith("NUMBER PAGE") or line.startswith("Spin Master"):
            continue
        description_lines.append(line)
    return _compact(" ".join(description_lines[:3]))


def _first_match(pattern: str, text: str, flags: int = 0) -> str:
    match = re.search(pattern, text, flags)
    if not match:
        return ""
    if "value" in match.groupdict():
        return match.group("value")
    return match.group(1)


def _compact(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _format_us_date(value: str) -> str:
    if not value:
        return ""
    try:
        return datetime.strptime(value, "%m/%d/%Y").date().isoformat()
    except ValueError:
        return value


def validate_pdf_against_schedule(rows: list[dict[str, Any]], pdf_data: dict[str, Any]) -> dict[str, Any]:
    po_number = str(pdf_data.get("po_number") or "")
    excel_rows = [row for row in rows if str(row.get("contract_base") or "") == po_number]
    by_line = {int(row["contract_line"]): row for row in excel_rows if row.get("contract_line") is not None}

    comparisons = []
    mismatch_count = 0
    for item in pdf_data.get("items", []):
        excel_row = by_line.get(int(item["line_no"]))
        if not excel_row:
            mismatch_count += 1
            comparisons.append(
                {
                    "line_no": item["line_no"],
                    "status": "missing",
                    "contract_no": "",
                    "checks": [],
                    "message": "Excel 中未找到对应合同号行",
                }
            )
            continue

        checks = [
            _compare("客PO", item.get("customer_po"), excel_row.get("customer_po"), kind="text"),
            _compare("数量", item.get("quantity"), excel_row.get("quantity"), kind="number"),
            _compare("Material Group", item.get("material_group"), excel_row.get("material_group"), kind="text"),
            _compare("Mat. Number", item.get("material_no"), excel_row.get("material_no"), kind="text"),
            _compare("Sales Material", item.get("sales_material"), excel_row.get("sales_material"), kind="text"),
            _compare("Sales Order", item.get("sales_order"), excel_row.get("sales_order"), kind="text"),
            _compare("Line Item", item.get("sales_line_item"), excel_row.get("sales_line_item"), kind="text"),
            _compare("DEL. DATE", item.get("delivery_date"), excel_row.get("po_ship_date"), kind="text"),
            _compare("NET VALUE", item.get("net_usd"), excel_row.get("total_usd"), kind="money"),
        ]
        row_mismatches = [check for check in checks if not check["ok"]]
        mismatch_count += len(row_mismatches)
        comparisons.append(
            {
                "line_no": item["line_no"],
                "status": "pass" if not row_mismatches else "warn",
                "contract_no": excel_row.get("contract_no"),
                "customer_po": excel_row.get("customer_po"),
                "excel_row": excel_row.get("excel_row"),
                "checks": checks,
                "message": "" if not row_mismatches else f"{len(row_mismatches)} 项不一致",
            }
        )

    pdf_quantity = float(pdf_data.get("total_quantity") or 0)
    excel_quantity = sum(float(row.get("quantity") or 0) for row in excel_rows)
    pdf_usd = float(pdf_data.get("total_net_usd") or 0)
    excel_usd = sum(float(row.get("total_usd") or 0) for row in excel_rows)

    return {
        "po_number": po_number,
        "pdf_item_count": len(pdf_data.get("items", [])),
        "excel_item_count": len(excel_rows),
        "matched_item_count": len([item for item in comparisons if item["status"] != "missing"]),
        "mismatch_count": mismatch_count,
        "status": "pass" if mismatch_count == 0 and len(excel_rows) == len(pdf_data.get("items", [])) else "warn",
        "totals": {
            "pdf_quantity": round(pdf_quantity, 2),
            "excel_quantity": round(excel_quantity, 2),
            "quantity_delta": round(excel_quantity - pdf_quantity, 2),
            "pdf_net_usd": round(pdf_usd, 2),
            "excel_net_usd": round(excel_usd, 2),
            "net_usd_delta": round(excel_usd - pdf_usd, 2),
        },
        "rows": comparisons,
    }


def _compare(label: str, pdf_value: Any, excel_value: Any, kind: str) -> dict[str, Any]:
    if kind == "number":
        pdf_number = float(pdf_value or 0)
        excel_number = float(excel_value or 0)
        ok = abs(pdf_number - excel_number) <= 0.01
        pdf_display: Any = round(pdf_number, 2)
        excel_display: Any = round(excel_number, 2)
    elif kind == "money":
        pdf_number = float(pdf_value or 0)
        excel_number = float(excel_value or 0)
        ok = abs(pdf_number - excel_number) <= 0.02
        pdf_display = round(pdf_number, 2)
        excel_display = round(excel_number, 2)
    else:
        pdf_display = _compact(pdf_value)
        excel_display = _compact(excel_value)
        ok = pdf_display.upper() == excel_display.upper()
    return {
        "label": label,
        "pdf": pdf_display,
        "excel": excel_display,
        "ok": ok,
    }


def export_rows_csv(rows: list[dict[str, Any]]) -> str:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([col.title for col in COLUMNS])
    for row in rows:
        writer.writerow([row.get(col.key, "") for col in COLUMNS])
    return output.getvalue()


def _empty_schedule(source: Path) -> dict[str, Any]:
    return {
        "sheet": SCHEDULE_SHEET,
        "source": str(source),
        "headers": [
            {"index": col.index, "key": col.key, "title": col.title, "kind": col.kind}
            for col in COLUMNS
        ],
        "rows": [],
        "stats": build_stats([]),
    }


def _empty_pdf(source: Path) -> dict[str, Any]:
    return {
        "source": str(source),
        "po_number": "",
        "order_date": "",
        "vendor": "",
        "buyer": "",
        "page_count": 0,
        "item_count": 0,
        "items": [],
        "total_quantity": 0,
        "total_net_usd": 0,
    }


def _empty_validation(pdf: dict[str, Any], schedule: dict[str, Any]) -> dict[str, Any]:
    return {
        "po_number": pdf.get("po_number", ""),
        "pdf_item_count": pdf.get("item_count", 0),
        "excel_item_count": 0,
        "matched_item_count": 0,
        "mismatch_count": 0,
        "status": "not_available",
        "totals": {
            "pdf_quantity": pdf.get("total_quantity", 0),
            "excel_quantity": 0,
            "quantity_delta": 0,
            "pdf_net_usd": pdf.get("total_net_usd", 0),
            "excel_net_usd": 0,
            "net_usd_delta": 0,
        },
        "rows": [],
        "schedule_row_count": schedule["stats"]["row_count"],
    }


def build_payload(excel_path: Path | str | None = None) -> dict[str, Any]:
    schedule_source = Path(excel_path or active_excel_path())
    schedule = load_schedule(excel_path=schedule_source) if schedule_source.exists() else _empty_schedule(schedule_source)

    pdf_source = active_pdf_path()
    pdf = parse_po_pdf(pdf_source) if pdf_source.exists() else _empty_pdf(pdf_source)
    validation = (
        validate_pdf_against_schedule(schedule["rows"], pdf)
        if schedule["rows"] and pdf["items"]
        else _empty_validation(pdf, schedule)
    )
    paths = workbook_paths()
    if excel_path is not None:
        paths["excel"] = str(schedule_source)
        paths["source"] = schedule_file_info(schedule_source)
    return {
        "paths": paths,
        "schedule": schedule,
        "pdf": pdf,
        "quality": run_quality_checks(schedule["rows"]),
        "validation": validation,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    }
