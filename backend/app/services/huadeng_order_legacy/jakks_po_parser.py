from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pdfplumber
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


def _clean(value: str) -> str:
    return re.sub(r"[ \t]+", " ", value or "").strip()


def _number(value: str) -> float:
    return float(re.sub(r"[^\d.-]", "", value.replace(",", "")) or 0)


def _product_name_from_description(value: str) -> str:
    """Prefer the PO's explicit Chinese name instead of copying a bilingual sentence."""
    description = _clean(str(value or "").replace("\n", " "))
    first_chinese = re.search(r"[\u3400-\u9fff]", description)
    if not first_chinese:
        return description
    chinese_name = description[first_chinese.start():]
    chinese_name = re.sub(r"\s+", "", chinese_name)
    return chinese_name.strip(" \t\r\n,，;；:：/") or description


def extract_text(path: str | Path) -> tuple[str, bool]:
    """Extract text, using OCR when the PDF is a scan without a text layer."""
    with pdfplumber.open(path) as pdf:
        text = "\n".join(page.extract_text(x_tolerance=2, y_tolerance=4) or "" for page in pdf.pages)
        if len(re.sub(r"\s+", "", text)) >= 80:
            return text, False

        rapidocr_texts: list[str] = []
        try:
            from app.services.carton_mark import get_rapidocr_engine, rapidocr_result_to_text

            rapidocr_engine = get_rapidocr_engine()
            if rapidocr_engine is not None:
                for page in pdf.pages:
                    image = page.to_image(resolution=180).original
                    recognized = rapidocr_result_to_text(rapidocr_engine(image))
                    if recognized.strip():
                        rapidocr_texts.append(recognized)
                rapidocr_text = "\n".join(rapidocr_texts)
                if len(re.sub(r"\s+", "", rapidocr_text)) >= 40:
                    return rapidocr_text, True
        except Exception:
            # A missing/failed RapidOCR runtime should not prevent the existing
            # Tesseract fallback from handling the scan.
            pass

        try:
            import pytesseract
        except ImportError as exc:
            if rapidocr_texts:
                return "\n".join(rapidocr_texts), True
            raise ValueError("扫描版 PO 需要服务器 OCR 组件（RapidOCR 或 Tesseract）") from exc
        pages = []
        for page in pdf.pages:
            image = page.to_image(resolution=300).original
            pages.append(pytesseract.image_to_string(image, lang="eng", config="--psm 6"))
        return "\n".join([*rapidocr_texts, *pages]), True


def _parse_text_fields(text: str) -> dict[str, Any]:
    compact = re.sub(r"[ \t]+", " ", text)
    contract = ""
    for pattern in (
        r"Contract\s*#?\s*[:：]\s*(\d{5,})",
        r"(?:No\.?|合同(?:编码|号码)?)[\s:：]*([A-Z]{1,12}[A-Z0-9-]*\d{5,})",
        r"\b([A-Z]{1,12}\d{6,})\b",
    ):
        found = re.search(pattern, compact, re.I)
        if found:
            contract = found.group(1).upper()
            break
    order_date = ""
    found = re.search(r"(?:Date|签定日期)[\s:：]*(\d{4})[./-](\d{1,2})[./-](\d{1,2})", compact, re.I)
    if found:
        order_date = f"{int(found.group(1)):04d}-{int(found.group(2)):02d}-{int(found.group(3)):02d}"
    if not order_date:
        found = re.search(
            r"Order Date\s*[:：]\s*(\d{1,2})[- /]([A-Za-z]{3,9})[- /](\d{4})",
            compact,
            re.I,
        )
        if found:
            for fmt in ("%d-%b-%Y", "%d-%B-%Y"):
                try:
                    order_date = datetime.strptime("-".join(found.groups()), fmt).date().isoformat()
                    break
                except ValueError:
                    pass
    if not order_date:
        found = re.search(
            r"(?<!Printed\s)\bDate\s*[:：]\s*(\d{1,2})[- /]([A-Za-z]{3,9})[- /](\d{4})",
            compact,
            re.I,
        )
        if found:
            for fmt in ("%d-%b-%Y", "%d-%B-%Y"):
                try:
                    order_date = datetime.strptime("-".join(found.groups()), fmt).date().isoformat()
                    break
                except ValueError:
                    pass
    ship_date = ""
    found = re.search(
        r"(?:Shipment date|装运期)[^\n\r]{0,30}?(?:BEFORE\s+)?([A-Za-z]+)[., ]+(\d{1,2})[., ]+(\d{4})",
        compact,
        re.I,
    )
    if found:
        for fmt in ("%B-%d-%Y", "%b-%d-%Y"):
            try:
                ship_date = datetime.strptime("-".join(found.groups()), fmt).date().isoformat()
                break
            except ValueError:
                pass

    customer = ""
    buyer_match = re.search(
        r"(?:The Buyers?|买方)\s*[:：]?\s*(?:[^\n\r]*[\n\r]+){0,3}"
        r"\s*([A-Z][A-Z0-9 .,&'()/-]+?(?:LIMITED|LTD\.?|INC\.?))\s*$",
        text,
        re.I | re.M,
    )
    if buyer_match:
        customer = _clean(buyer_match.group(1))
    contact = ""
    contact_match = re.search(
        r"(?:Attention|Attn)\s*[:：]\s*((?:Mr|Ms|Mrs)\.?\s+[A-Za-z][A-Za-z .'-]{1,40})",
        text,
        re.I,
    )
    if contact_match:
        contact = _clean(contact_match.group(1))
    return {
        "contract_no": contract,
        "order_date": order_date,
        "ship_date": ship_date,
        "customer": customer,
        "contact": contact,
    }


def _validate_lines(lines: list[dict[str, Any]]) -> list[str]:
    warnings = []
    if not lines:
        raise ValueError("PO 没有识别到产品明细行。")
    for line in lines:
        if line["unit"] in {"KGM", "KG"}:
            warnings.append(
                f"{line['item_no']}：数量按合同重量 {line['quantity']} {line['unit']} 写入，"
                "不换算为PCS；单位保留在“识别明细”，不自动写入“特别备注”。"
            )
        expected = round(line["quantity"] * line["unit_price_usd"], 2)
        if abs(expected - line["total_usd"]) > max(0.05, abs(line["total_usd"]) * 0.001):
            warnings.append(
                f"{line['item_no']}：数量×单价={expected}，与金额 {line['total_usd']} 不一致。"
            )
    return warnings


def _parse_pdf_po(path: Path) -> dict[str, Any]:
    text, used_ocr = extract_text(path)
    fields = _parse_text_fields(text)
    lines = []
    row_pattern = re.compile(
        r"(?m)^\s*([A-Z0-9]+(?:-[A-Z0-9]+){1,6})\s+(.+?)\s+"
        r"([\d,.]+)\s+(KGM|KG|PCE|PCS)\s+US?[$S]\s*([\d,.]+)\s+US?[$S]\s*([\d,.]+)\s*$",
        re.I,
    )
    for match in row_pattern.finditer(text):
        qty = _number(match.group(3))
        unit_price = _number(match.group(5))
        amount = _number(match.group(6))
        description = _clean(match.group(2))
        lines.append({
            "item_no": match.group(1).upper(),
            "product_name": _product_name_from_description(description),
            "po_description": description,
            "product_name_source": "PO中文品名",
            "quantity": qty,
            "unit": match.group(4).upper(),
            "unit_price_usd": unit_price,
            "total_usd": amount,
            "special_note": "",
        })
    if not lines:
        # OCR often breaks one table row into two lines.  Join adjacent lines and retry.
        joined = re.sub(r"\n(?!(?:[A-Z0-9]+(?:-[A-Z0-9]+){1,6})\b)", " ", text)
        for match in row_pattern.finditer(joined):
            description = _clean(match.group(2))
            lines.append({
                "item_no": match.group(1).upper(),
                "product_name": _product_name_from_description(description),
                "po_description": description,
                "product_name_source": "PO中文品名",
                "quantity": _number(match.group(3)),
                "unit": match.group(4).upper(),
                "unit_price_usd": _number(match.group(5)),
                "total_usd": _number(match.group(6)),
                "special_note": "",
            })
    warnings = []
    if used_ocr:
        warnings.append("扫描版 PDF 已使用 OCR；合同号、货号、数量和金额需做算术复核。")
    if not fields["contract_no"]:
        warnings.append("未识别合同号。")
    warnings.extend(_validate_lines(lines))
    return {
        "filename": path.name,
        **fields,
        "lines": lines,
        "warnings": list(dict.fromkeys(warnings)),
        "used_ocr": used_ocr,
        "source_format": "pdf",
    }


def _header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", str(value or "").lower())


def _find_header_columns(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    aliases = {
        "item_no": ("itemno", "型号", "貨號", "货号"),
        "product_name": ("namesanddescription", "namesanddescriptions", "品名描述", "產品名稱", "产品名称"),
        "quantity": ("quantities", "quantity", "數量", "数量"),
        "unit": ("unit", "單位", "单位"),
        "unit_price_usd": ("unitprice", "單價", "单价"),
        "total_usd": ("amount", "金額", "金额"),
        "special_note": ("remarks", "remark", "notes", "note", "备注", "借註", "備註"),
    }
    required = {
        "item_no",
        "product_name",
        "quantity",
        "unit",
        "unit_price_usd",
        "total_usd",
    }
    for row_index, row in enumerate(rows):
        mapped: dict[str, int] = {}
        for column_index, value in enumerate(row):
            key = _header_key(value)
            if not key:
                continue
            for field, candidates in aliases.items():
                if field not in mapped and any(candidate in key for candidate in candidates):
                    mapped[field] = column_index
        if required.issubset(mapped):
            return row_index, mapped
    raise ValueError("WPS 转换 Excel 中未找到 Jakks 产品明细表头。")


def _excel_date(value: Any, epoch: datetime) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and 30000 <= value <= 80000:
        return from_excel(value, epoch).date().isoformat()
    match = re.search(r"(\d{4})[./-](\d{1,2})[./-](\d{1,2})", str(value or ""))
    if match:
        return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    return ""


def _english_date(value: Any) -> str:
    text = _clean(str(value or ""))
    match = re.search(r"(\d{1,2})[- /]([A-Za-z]{3,9})[- /](\d{4})", text)
    if not match:
        return ""
    for fmt in ("%d-%b-%Y", "%d-%B-%Y"):
        try:
            return datetime.strptime("-".join(match.groups()), fmt).date().isoformat()
        except ValueError:
            pass
    return ""


def _value_after_label(rows: list[list[Any]], pattern: str) -> str:
    label = re.compile(pattern, re.I)
    for row in rows:
        for index, value in enumerate(row):
            raw_text = str(value or "").strip()
            text = _clean(raw_text.replace("\n", " "))
            match = label.search(text)
            if not match:
                continue
            inline = text[match.end():].strip(" :：")
            if inline:
                return inline
            for candidate in row[index + 1:]:
                candidate_raw = str(candidate or "").strip()
                candidate_text = _clean(candidate_raw.replace("\n", " "))
                if candidate_text and candidate_text not in {"0", "0.0"}:
                    return candidate_raw
    return ""


def _display_cell_value(cell: Any) -> Any:
    """Preserve identifier zeros represented by an Excel number format."""
    value = cell.value
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return value
    if float(value).is_integer():
        number_format = str(cell.number_format or "").split(";", 1)[0]
        simplified = re.sub(r'"[^"]*"', "", number_format)
        simplified = re.sub(r"_.", "", simplified)
        simplified = re.sub(r"\\.", "", simplified)
        simplified = re.sub(r"\*.", "", simplified)
        simplified = re.sub(r"\[[^\]]*\]", "", simplified)
        simplified = re.sub(r"\s+", "", simplified)
        if re.fullmatch(r"0+", simplified):
            return f"{int(value):0{len(simplified)}d}"
    return value


def _standard_contract_fields(rows: list[list[Any]], text: str) -> dict[str, Any]:
    fields = _parse_text_fields(text)
    customer_po = _value_after_label(rows, r"Customer\s*PO\s*[:：]?")
    confirmation_text = _value_after_label(rows, r"Confirmation\s*No\.?\s*[:：]?")
    paired_customer = ""
    if re.search(r"Confirmation\s*No|Port\s*of\s*Loading", customer_po, re.I):
        for row in rows:
            nonempty = [str(value).strip() for value in row if value not in (None, "")]
            for index, value in enumerate(nonempty[:-1]):
                if "Customer PO:" not in value or "Confirmation No:" not in value:
                    continue
                values = [_clean(line) for line in nonempty[index + 1].splitlines() if _clean(line)]
                if len(values) >= 3:
                    customer_po = values[1]
                    confirmation_text = values[2]
                    if len(values) >= 6:
                        paired_customer = values[5]
                break
    confirmation_match = re.search(r"\b(\d{5,})\b(?:\s+(.+))?", confirmation_text)
    confirmation_no = confirmation_match.group(1) if confirmation_match else ""
    country = _clean(confirmation_match.group(2) or "") if confirmation_match else ""

    consignee = _value_after_label(rows, r"Ultimate\s*Consignee\s*Name\s*[:：]?")
    consignee_lines = [
        _clean(line)
        for line in re.split(r"[\r\n]+", consignee)
        if _clean(line) and not re.fullmatch(r"\d+", _clean(line))
    ]
    customer = paired_customer or (consignee_lines[0] if consignee_lines else "")
    if customer and country and country.upper() not in customer.upper():
        customer = f"{customer} {country}"

    fields.update({
        "customer_po": customer_po,
        "confirmation_no": confirmation_no,
        "country": country,
    })
    if customer:
        fields["customer"] = customer
    return fields


def _parse_standard_contract_rows(
    rows: list[list[Any]],
) -> list[dict[str, Any]]:
    """Parse JAKKS' normal CONTRACT layout after WPS PDF→Excel conversion."""
    lines: list[dict[str, Any]] = []
    item_pattern = re.compile(r"^([A-Z0-9]{5,}(?:-[A-Z0-9]+){0,6})\b", re.I)
    for row_index, row in enumerate(rows):
        item = ""
        item_column = -1
        for column_index, value in enumerate(row):
            cell_text = str(value or "").strip()
            match = item_pattern.match(cell_text)
            if match and re.search(r"(?:UPC|Stock\s*#|Country\s+of\s+Origin)", cell_text, re.I):
                item = match.group(1).upper()
                item_column = column_index
                break
        if not item:
            continue

        description = ""
        for value in row[item_column + 1:]:
            candidate = _clean(str(value or "").replace("\n", " "))
            if re.search(r"Customer\s*Item\s*No", candidate, re.I):
                description = re.split(r"Customer\s*Item\s*No\.?\s*:", candidate, flags=re.I)[0].strip()
                break
        if not description:
            candidates = [
                _clean(str(value or "").replace("\n", " "))
                for value in row[item_column + 1:]
                if re.search(r"[A-Za-z]{4}", str(value or ""))
                and not re.search(r"(?:Outer\s*Pack|USD|Request\s*Date)", str(value or ""), re.I)
            ]
            description = candidates[0] if candidates else ""

        quantity = 0.0
        cartons = 0.0
        outer_pack = 0.0
        quantity_column = -1
        for column_index, value in enumerate(row):
            block = str(value or "")
            if not re.search(r"(?:\bCA\b|Outer\s*Pack)", block, re.I):
                continue
            first_number = re.search(r"^\s*([\d,.]+)", block)
            carton_match = re.search(r"([\d,.]+)\s*CA\b", block, re.I)
            outer_match = re.search(r"Outer\s*Pack\s*:\s*([\d,.]+)", block, re.I)
            quantity = _number(first_number.group(1)) if first_number else 0
            cartons = _number(carton_match.group(1)) if carton_match else 0
            outer_pack = _number(outer_match.group(1)) if outer_match else 0
            quantity_column = column_index
            break

        amount = 0.0
        amount_column = -1
        for column_index, value in enumerate(row):
            if re.search(r"\bUSD\b", str(value or ""), re.I):
                amount = _number(str(value))
                amount_column = column_index
                break

        unit_price = 0.0
        if amount_column > 0:
            for column_index in range(amount_column - 1, item_column, -1):
                if column_index == quantity_column:
                    continue
                value = row[column_index]
                if isinstance(value, (int, float)) and 0 < float(value) < 10000:
                    unit_price = float(value)
                    break
                candidate = _clean(str(value or ""))
                if re.fullmatch(r"\d+(?:\.\d+)?", candidate):
                    unit_price = float(candidate)
                    break
        if not unit_price and amount and quantity:
            unit_price = amount / quantity

        ship_date = ""
        special_note = ""
        for following in rows[row_index: min(len(rows), row_index + 8)]:
            joined = " ".join(
                _clean(str(value or "").replace("\n", " "))
                for value in following
                if value not in (None, "")
            )
            if not ship_date and re.search(r"Request\s*Date", joined, re.I):
                ship_date = _english_date(joined)
            # JAKKS has both ``ORDER NOTES: REMARKS`` and the business
            # ``NOTES:`` field.  Only the latter belongs in the schedule's
            # special-note column.
            note_match = re.search(r"(?<!ORDER )\bNOTES\s*:\s*(.+)", joined, re.I)
            if note_match and not special_note:
                special_note = _clean(note_match.group(1))
            if following is not row and any(
                item_pattern.match(str(value or "").strip())
                and re.search(r"(?:UPC|Stock\s*#|Country\s+of\s+Origin)", str(value or ""), re.I)
                for value in following
            ):
                break

        if not description or quantity <= 0 or amount <= 0:
            continue
        lines.append({
            "item_no": item,
            "product_name": description,
            "po_description": description,
            "product_name_source": "PO英文品名",
            "quantity": quantity,
            "unit": "PCS",
            "inner_pack": "",
            "outer_pack": outer_pack or "",
            "cartons": cartons or "",
            "unit_price_usd": unit_price,
            "total_usd": amount,
            "ship_date": ship_date,
            "special_note": special_note,
        })
    return lines


def _parse_xlsx_po(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        rows: list[list[Any]] = []
        for worksheet in workbook.worksheets:
            rows.extend(
                [
                    [_display_cell_value(cell) for cell in row]
                    for row in worksheet.iter_rows()
                ]
            )
        if not rows:
            raise ValueError("WPS 转换 Excel 没有可读取内容。")
        text = "\n".join(
            " ".join(str(value) for value in row if value not in (None, ""))
            for row in rows
        )
        is_standard_contract = (
            re.search(r"JAKKS\s+PACIFIC", text, re.I)
            and re.search(r"\bCONTRACT\b", text, re.I)
            and re.search(r"ITEM\s+NUMBER", text, re.I)
        )
        lines = []
        if is_standard_contract:
            fields = _standard_contract_fields(rows, text)
            lines = _parse_standard_contract_rows(rows)
        else:
            header_row, columns = _find_header_columns(rows)
            item_pattern = re.compile(r"^[A-Z0-9]+(?:-[A-Z0-9]+){1,6}$", re.I)
            for row in rows[header_row + 1:]:
                item = _clean(str(row[columns["item_no"]] or "")).upper()
                if not item_pattern.fullmatch(item):
                    continue
                quantity = _number(str(row[columns["quantity"]] or ""))
                unit_price = _number(str(row[columns["unit_price_usd"]] or ""))
                amount = _number(str(row[columns["total_usd"]] or ""))
                unit = _clean(str(row[columns["unit"]] or "")).upper()
                description = _clean(
                    str(row[columns["product_name"]] or "").replace("\n", " ")
                )
                product_name = _product_name_from_description(description)
                special_note = (
                    _clean(str(row[columns["special_note"]] or "").replace("\n", " "))
                    if "special_note" in columns
                    else ""
                )
                if not product_name or quantity <= 0 or not unit:
                    continue
                lines.append({
                    "item_no": item,
                    "product_name": product_name,
                    "po_description": description,
                    "product_name_source": "PO中文品名",
                    "quantity": quantity,
                    "unit": unit,
                    "unit_price_usd": unit_price,
                    "total_usd": amount,
                    "special_note": special_note,
                })
            fields = _parse_text_fields(text)

        if not fields["order_date"]:
            for row in rows:
                joined = " ".join(str(value or "") for value in row)
                if re.search(r"(?:Date|签定日期)", joined, re.I):
                    for value in row:
                        parsed = _excel_date(value, workbook.epoch)
                        if parsed:
                            fields["order_date"] = parsed
                            break
                if fields["order_date"]:
                    break
        if not fields["customer"] and re.search(
            r"ROYAL\s+REGENT\s+PRODUCTS\s+INDUSTRIES\s+(?:LIMITED|LTD\.?)",
            text,
            re.I,
        ):
            fields["customer"] = "ROYAL REGENT PRODUCTS INDUSTRIES LIMITED"

        warnings = [
            (
                "已识别 Jakks 标准 CONTRACT / WPS Excel；关键字段已按标签与产品块定位并执行数量×单价复核。"
                if is_standard_contract
                else "已读取 WPS PDF转Excel 结果；关键字段已按表头定位并执行数量×单价算术复核。"
            )
        ]
        if not fields["contract_no"]:
            warnings.append("未识别合同号。")
        warnings.extend(_validate_lines(lines))
        return {
            "filename": path.name,
            **fields,
            "lines": lines,
            "warnings": list(dict.fromkeys(warnings)),
            "used_ocr": True,
            "source_format": "wps_excel",
        }
    finally:
        workbook.close()


def parse_po(path: str | Path) -> dict[str, Any]:
    po_path = Path(path)
    if re.search(r"(?:^|[-_])(CXL|SUP)(?:[-_]|$)", po_path.stem, re.I):
        raise ValueError(
            "识别为取消/补充/修改单（文件名含 CXL/SUP）；按业务规则只处理新单，"
            "本文件已拦截，不会生成 Excel 或写入数据库。"
        )
    suffix = po_path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf_po(po_path)
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_xlsx_po(po_path)
    raise ValueError("Jakks PO 仅支持 PDF、WPS转换后的 XLSX 或 XLSM。")
