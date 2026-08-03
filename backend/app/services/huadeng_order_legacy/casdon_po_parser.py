# -*- coding: utf-8 -*-
"""Casdon PO parser.

Supports the Casdon PDF purchase orders in the current sample set and Excel
files produced from the same layout. The output shape is intentionally close
to the ZURU system's parser: one header plus a list of normalized line items.
"""
from __future__ import annotations

import os
import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, asdict
from datetime import datetime
from typing import Iterable

import openpyxl
import pdfplumber


def normalize_date(value) -> str:
    """Return YYYY-MM-DD for common Casdon date formats."""
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip().replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text)
    # dd/mm/yyyy or dd-mm-yyyy
    m = re.match(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", text)
    if m:
        day, month, year = int(m.group(1)), int(m.group(2)), m.group(3)
        return f"{year}-{month:02d}-{day:02d}"
    # yyyy/mm/dd or yyyy-mm-dd
    m = re.match(r"(\d{4})[/-](\d{1,2})[/-](\d{1,2})", text)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    return text


def _num(value: str) -> float:
    if value is None:
        return 0.0
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


def _int(value: str) -> int:
    return int(round(_num(value)))


def _clean_excel_text(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ")
    text = text.replace("\r", "\n").replace("\n", " ")
    text = re.sub(r"\s*/\s*", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _excel_header_text(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ")
    text = text.replace("\r", "\n").replace("\n", " / ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def _parse_qty_unit(qty_cell: str, unit_cells) -> tuple[int, str]:
    qty = 0
    unit = ""
    m = re.search(r"([\d,]+)\s*([A-Za-z]+)?", qty_cell or "")
    if m:
        qty = _int(m.group(1))
        unit = (m.group(2) or "").strip()
    if not unit:
        for cell in unit_cells:
            text = _clean_excel_text(cell.value)
            if re.fullmatch(r"[A-Za-z]+", text):
                unit = text
                break
    return qty, unit


def _last_two_numbers(cells) -> tuple[float, float]:
    values = []
    for cell in cells:
        text = _clean_excel_text(cell.value)
        if not text:
            continue
        if re.fullmatch(r"[\d,]+(?:\.\d+)?", text):
            values.append(_num(text))
    if len(values) >= 2:
        return values[-2], values[-1]
    if len(values) == 1:
        return values[0], 0.0
    return 0.0, 0.0


def _field(pattern: str, text: str, default: str = "") -> str:
    m = re.search(pattern, text, re.IGNORECASE)
    return re.sub(r"\s+", " ", m.group(1)).strip() if m else default


@dataclass
class CasdonLine:
    line_no: str
    item_code: str
    description_en: str
    qty: int = 0
    unit: str = ""
    unit_price: float = 0.0
    total_usd: float = 0.0
    commodity_code: str = ""
    is_charge: bool = False

    def to_dict(self) -> dict:
        return asdict(self)


class CasdonPOParser:
    PRODUCT_RE = re.compile(
        r"^(\d+)\s+([A-Za-z]*\d[\w./-]*)\s+(.+?)\s+(\d{8})\s+"
        r"([\d,]+)\s+([A-Za-z]+)\s+([\d,]+(?:\.\d+)?)\s+([\d,]+(?:\.\d+)?)$"
    )
    CHARGE_RE = re.compile(
        r"^(\d+)\s+(.+?(?:Charge|Charges|Freight|Container).+?)\s+"
        r"([\d,]+)\s+([A-Za-z]+)\s+([\d,]+(?:\.\d+)?)\s+([\d,]+(?:\.\d+)?)$",
        re.IGNORECASE,
    )

    def parse(self, path: str) -> dict:
        ext = os.path.splitext(path)[1].lower()
        if ext == ".pdf":
            text = self._text_from_pdf(path)
        elif ext in (".xlsx", ".xlsm"):
            structured = self._parse_excel_workbook(path)
            if structured and structured.get("lines"):
                return structured
            text = self._text_from_excel(path)
        else:
            raise ValueError(f"不支持的文件格式: {ext}")

        header = self._parse_header(text)
        lines = [line.to_dict() for line in self._parse_lines(text)]
        if not lines:
            raise ValueError("未识别到产品行")
        return {**header, "lines": lines, "raw_text": text[:8000]}

    def _text_from_pdf(self, path: str) -> str:
        parts = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                parts.append(page.extract_text(x_tolerance=1, y_tolerance=3) or "")
        return "\n".join(parts)

    def _text_from_excel(self, path: str) -> str:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            rows = []
            for row in ws.iter_rows(max_row=1000, max_col=80):
                vals = []
                for cell in row:
                    if cell.value is not None:
                        vals.append(_clean_excel_text(cell.value))
                if vals:
                    rows.append(" ".join(vals))
            drawing_text = self._drawing_text_from_excel(path)
            if drawing_text:
                rows.append(drawing_text)
            return "\n".join(rows)
        finally:
            wb.close()

    def _parse_excel_workbook(self, path: str) -> dict:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            header = self._parse_excel_header(ws)
            if str(header.get("your_reference") or "").lower().startswith("all amounts"):
                header["your_reference"] = ""
            lines = [line.to_dict() for line in self._parse_excel_lines(ws)]
            if not lines:
                return {}
            text = self._text_from_excel(path)
            drawing_header = self._parse_drawing_header(path)
            text_header = self._parse_header(text)
            for key in (
                "po_number",
                "po_date",
                "ship_date",
                "your_reference",
                "account_number",
                "version",
                "customer",
                "customer_po_header",
                "container_type",
            ):
                if not header.get(key):
                    header[key] = drawing_header.get(key) or text_header.get(key) or ""
            if not header.get("po_number"):
                header["po_number"] = self._po_from_filename(path)
            if not header.get("version"):
                header["version"] = self._parse_version(text)
            return {**header, "lines": lines, "raw_text": text[:8000]}
        finally:
            wb.close()

    @staticmethod
    def _drawing_shapes_from_excel(path: str) -> list[str]:
        """Extract text boxes kept by WPS when PDF fields are not real worksheet cells."""
        shapes: list[str] = []
        try:
            with zipfile.ZipFile(path) as archive:
                names = sorted(
                    name
                    for name in archive.namelist()
                    if name.startswith("xl/drawings/drawing") and name.endswith(".xml")
                )
                for name in names:
                    root = ET.fromstring(archive.read(name))
                    text_tag = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
                    for anchor in list(root):
                        parts = [node.text or "" for node in anchor.iter(text_tag)]
                        value = re.sub(r"\s+", " ", "".join(parts)).strip()
                        if value:
                            shapes.append(value)
        except (OSError, zipfile.BadZipFile, ET.ParseError):
            return []
        return shapes

    def _drawing_text_from_excel(self, path: str) -> str:
        return "\n".join(self._drawing_shapes_from_excel(path))

    def _parse_drawing_header(self, path: str) -> dict:
        shapes = self._drawing_shapes_from_excel(path)
        text = "\n".join(shapes)
        header_shape = next(
            (
                shape
                for shape in shapes
                if "Purchase Order No" in shape
                or (
                    re.search(r"(?<!\d)\d{7,12}(?!\d)", shape)
                    and re.search(r"\d{1,2}[/-]\d{1,2}[/-]\d{4}", shape)
                )
            ),
            "",
        )
        po = ""
        dates: list[str] = []
        ref = ""
        account = ""
        if header_shape:
            po_match = re.search(r"(?<!\d)(\d{7,12})(?!\d)", header_shape)
            po = po_match.group(1) if po_match else ""
            dates = [
                normalize_date(match.group(0))
                for match in re.finditer(r"\d{1,2}[/-]\d{1,2}[/-]\d{4}", header_shape)
            ]
            account_match = re.search(r"(?<![A-Z0-9])(ROY)\s*(\d{2,5})(?!\d)", header_shape, re.I)
            if account_match:
                account = f"{account_match.group(1)}{account_match.group(2)}".upper()
            if po:
                ref_match = re.search(
                    rf"(?<!\d){re.escape(po)}(?!\d)(.*?)(?:US\s*Dollar|All\s*Amounts)",
                    header_shape,
                    re.I,
                )
                if ref_match:
                    ref = re.sub(r"\bVARIOUS\b", "", ref_match.group(1), flags=re.I).strip()

        customer, cpo = self._parse_reference(ref)
        if not customer:
            deliver = _field(r"Deliver\s+To\s*([^\n]+)", text)
            customer = self._customer_from_deliver_to(deliver)
        return {
            "po_number": po,
            "po_date": dates[0] if len(dates) >= 1 else "",
            "ship_date": dates[1] if len(dates) >= 2 else "",
            "your_reference": ref,
            "account_number": account,
            "version": self._parse_version(text),
            "customer": customer,
            "customer_po_header": cpo,
            "container_type": self._parse_container_type(ref),
        }

    @staticmethod
    def _po_from_filename(path: str) -> str:
        name = os.path.basename(path)
        match = re.search(r"(?:Purchase\s+Order|PO)[^\d]*(\d{7,12})", name, re.I)
        return match.group(1) if match else ""

    @staticmethod
    def _customer_from_deliver_to(value: str) -> str:
        text = re.sub(r"\s+", " ", str(value or "")).strip().upper()
        known = (
            "TARGET",
            "WALMART",
            "AMAZON",
            "COSTCO",
            "WHB",
            "KIDOODLE",
            "MODERN BRANDS",
            "TRENDY WORLD",
        )
        for customer in known:
            if customer in text:
                return customer
        return ""

    def _parse_excel_header(self, ws) -> dict:
        value_blob = ""
        label_row = 0
        for row in ws.iter_rows(min_row=1, max_row=40, max_col=25):
            for cell in row:
                text = _clean_excel_text(cell.value)
                if "Purchase Order No" not in text:
                    continue
                label_row = cell.row
                for next_cell in row[cell.column:]:
                    next_text = _excel_header_text(next_cell.value)
                    if re.search(r"\d{6,}", next_text):
                        value_blob = next_text
                        break
                break
            if value_blob:
                break

        tokens = [p.strip() for p in re.split(r"\s+/\s+", value_blob) if p and p.strip()]
        po = next((t for t in tokens if re.fullmatch(r"\d{6,}", t)), "")
        dates = [normalize_date(t) for t in tokens if re.match(r"\d{1,2}[/-]\d{1,2}[/-]\d{4}$", t)]
        po_date = dates[0] if len(dates) >= 1 else ""
        ship_date = dates[1] if len(dates) >= 2 else ""
        account = next((t for t in tokens if re.fullmatch(r"[A-Z]{2,}\d+", t, re.I)), "")
        ref_parts = []
        if po:
            after_po = False
            for token in tokens:
                if token == po:
                    after_po = True
                    continue
                if not after_po:
                    continue
                if token.upper() == "US DOLLAR" or "days from document" in token.lower():
                    break
                ref_parts.append(token)
        ref = " ".join(ref_parts).strip()
        # The converter often puts VARIOUS as a separate visual line under reference.
        ref = re.sub(r"\bVARIOUS\b", "", ref, flags=re.I).strip()
        ref = re.split(r"\bUS\s+Dollar\b|days\s+from\s+document", ref, flags=re.I)[0]
        ref = ref.replace("/", " ").strip()

        full_text = ""
        if label_row:
            fragments = []
            for row in ws.iter_rows(min_row=1, max_row=80, max_col=25):
                for cell in row:
                    text = _clean_excel_text(cell.value)
                    if text:
                        fragments.append(text)
            full_text = "\n".join(fragments)

        version = self._parse_version(full_text)
        customer, cpo = self._parse_reference(ref)
        return {
            "po_number": po,
            "po_date": po_date,
            "ship_date": ship_date,
            "your_reference": ref,
            "account_number": account,
            "version": version,
            "customer": customer,
            "customer_po_header": cpo,
            "container_type": self._parse_container_type(ref),
        }

    def _parse_excel_lines(self, ws) -> list[CasdonLine]:
        header_row = 0
        header_cells = None
        for row in ws.iter_rows(min_row=1, max_row=80, max_col=25):
            row_text = " ".join(_clean_excel_text(cell.value) for cell in row)
            if "Line" in row_text and "Item Code" in row_text and "Description" in row_text:
                header_row = row[0].row
                header_cells = row
                break
        if not header_row:
            return []

        columns: dict[str, int] = {}
        header_aliases = {
            "line": ("line",),
            "item": ("item code",),
            "description": ("description",),
            "commodity": ("commodity code",),
            "qty": ("quantity",),
            "unit": ("unit",),
            "price": ("unit price",),
            "net": ("net",),
        }
        for cell in header_cells or ():
            label = re.sub(r"\s+", " ", _clean_excel_text(cell.value)).lower()
            for key, aliases in header_aliases.items():
                if key in columns:
                    continue
                if any(alias == label or alias in label for alias in aliases):
                    # "Unit Price" belongs to price, not the bare "Unit" column.
                    if key == "unit" and "price" in label:
                        continue
                    columns[key] = cell.column

        line_col = columns.get("line", 1)
        item_col = columns.get("item", 2)
        desc_col = columns.get("description", 3)
        commodity_col = columns.get("commodity", 4)
        qty_col = columns.get("qty")
        unit_col = columns.get("unit")
        price_col = columns.get("price")
        net_col = columns.get("net")

        lines: list[CasdonLine] = []
        for row in ws.iter_rows(min_row=header_row + 1, max_row=header_row + 120, max_col=25):
            line_text = _clean_excel_text(row[line_col - 1].value)
            combined = re.match(r"^(\d+)\s+([A-Za-z]*\d[\w./-]*)$", line_text)
            line_no = combined.group(1) if combined else line_text
            if not re.fullmatch(r"\d+", line_no):
                continue
            if combined:
                item = combined.group(2)
                actual_item_col = line_col
            else:
                item = _clean_excel_text(row[item_col - 1].value) if item_col != line_col else ""
                actual_item_col = item_col
                if not item:
                    for col in range(line_col + 1, min(line_col + 4, len(row)) + 1):
                        candidate = _clean_excel_text(row[col - 1].value)
                        if re.match(r"[A-Za-z]*\d", candidate) or re.search(
                            r"(Charge|Charges|Freight|Container)", candidate, re.I
                        ):
                            item = candidate
                            actual_item_col = col
                            break

            desc = _clean_excel_text(row[desc_col - 1].value) if desc_col not in (line_col, actual_item_col) else ""
            actual_desc_col = desc_col
            if not desc:
                for col in range(actual_item_col + 1, min(actual_item_col + 4, len(row)) + 1):
                    candidate = _clean_excel_text(row[col - 1].value)
                    if candidate and not re.fullmatch(r"\d{8}", candidate):
                        desc = candidate
                        actual_desc_col = col
                        break

            commodity = (
                _clean_excel_text(row[commodity_col - 1].value)
                if commodity_col not in (line_col, actual_item_col, actual_desc_col)
                else ""
            )
            actual_commodity_col = commodity_col
            if not re.search(r"\d{8}", commodity):
                for col in range(actual_desc_col + 1, min(actual_desc_col + 5, len(row)) + 1):
                    candidate = _clean_excel_text(row[col - 1].value)
                    if re.match(r"^\d{8}(?:\s+[\d,]+)?$", candidate):
                        commodity = candidate
                        actual_commodity_col = col
                        break

            qty_text = _clean_excel_text(row[qty_col - 1].value) if qty_col else ""
            is_charge_item = bool(re.search(r"(Charge|Charges|Freight|Container)", item, re.I))
            if not qty_text and is_charge_item and re.fullmatch(r"[\d,]+(?:\.0+)?", commodity):
                qty_text = commodity
            if not qty_text:
                embedded = re.search(r"\b\d{8}\s+([\d,]+)\s*$", commodity)
                qty_text = embedded.group(1) if embedded else ""
            if not qty_text:
                for col in range(actual_commodity_col + 1, min(actual_commodity_col + 8, len(row)) + 1):
                    candidate = _clean_excel_text(row[col - 1].value)
                    if re.fullmatch(r"[\d,]+", candidate) and _int(candidate) > 0:
                        qty_text = candidate
                        break
            qty = _int(qty_text)
            unit = _clean_excel_text(row[unit_col - 1].value) if unit_col else ""
            if not unit:
                _, unit = _parse_qty_unit(qty_text, row[max(actual_commodity_col, 1):16])

            price = _num(row[price_col - 1].value) if price_col else 0.0
            net = _num(row[net_col - 1].value) if net_col else 0.0
            if price <= 0:
                price, fallback_net = _last_two_numbers(row[7:20])
                net = net or fallback_net
            if net <= 0 and price_col:
                trailing = [
                    _num(cell.value)
                    for cell in row[price_col:20]
                    if re.fullmatch(r"[\d,]+(?:\.\d+)?", _clean_excel_text(cell.value))
                ]
                if trailing:
                    net = trailing[-1]
            if net <= 0 and qty > 0 and price > 0:
                net = round(qty * price, 2)
            if qty <= 0 or (price <= 0 and net <= 0):
                continue

            if is_charge_item:
                lines.append(CasdonLine(
                    line_no=line_no,
                    item_code="额外运费",
                    description_en=f"{item} {desc}".strip(),
                    qty=qty,
                    unit=unit,
                    unit_price=price,
                    total_usd=net,
                    is_charge=True,
                ))
                continue

            if not item or not re.match(r"[A-Za-z]*\d", item):
                continue

            lines.append(CasdonLine(
                line_no=line_no,
                item_code=item,
                description_en=desc,
                commodity_code=re.sub(r"\s+[\d,]+\s*$", "", commodity),
                qty=qty,
                unit=unit,
                unit_price=price,
                total_usd=net,
            ))
        return lines

    def _parse_header(self, text: str) -> dict:
        po = _field(r"Purchase\s+Order\s+No:\s*(\d+)", text)
        ref = _field(r"Your\s+Reference\s+(.+?)(?:\n|Royal\s+Regent|All\s+Amounts)", text)
        if "[" in ref and "]" not in ref:
            # Long references wrap onto the next visual line, e.g.
            # "DAM [SHOPPERS DRUG" / "... MART] 9114023". Stitch the closing part back.
            ref = self._complete_wrapped_reference(ref, text)
        order_date = normalize_date(_field(r"Order\s+Date\s+(\d{1,2}[/-]\d{1,2}[/-]\d{4})", text))
        ship_date = normalize_date(_field(r"Shipment\s+Date\s+(\d{1,2}[/-]\d{1,2}[/-]\d{4})", text))
        account = _field(r"Account\s+Number\s+([A-Z0-9-]+)", text)
        version = self._parse_version(text)
        customer, cpo = self._parse_reference(ref)
        return {
            "po_number": po,
            "po_date": order_date,
            "ship_date": ship_date,
            "your_reference": ref,
            "account_number": account,
            "version": version,
            "customer": customer,
            "customer_po_header": cpo,
            "container_type": self._parse_container_type(ref),
        }

    @staticmethod
    def _complete_wrapped_reference(ref: str, text: str) -> str:
        """Recover a bracketed customer name whose closing part wrapped to the next line.

        The wrapped fragment lands inside the supplier-address line (e.g.
        "Royal Regent Products (H.K) Ltd MART] 9114023"), so we only take the
        trailing ALL-CAPS tokens right before the closing bracket.
        """
        lines = text.splitlines()
        for i, line in enumerate(lines):
            if "Your Reference" not in line:
                continue
            for follow in lines[i + 1 : i + 4]:
                m = re.search(r"((?:[A-Z][A-Z0-9&'.-]*\s+)*[A-Z][A-Z0-9&'.-]*)\]", follow)
                if m:
                    stitched = f"{ref} {m.group(1)}]"
                    # Text after "]" on the wrapped line (e.g. customer PO numbers
                    # "9114023 / 9114040 LCL") still belongs to the reference.
                    remainder = follow[m.end():].strip()
                    if remainder:
                        stitched = f"{stitched} {remainder}"
                    return stitched
            break
        return ref

    @staticmethod
    def _parse_version(text: str) -> str:
        patterns = [
            r"\b(AW\d{2})\s+version\b",
            r"\bunder\s+(AW\d{2})\s+version\b",
            r"\bAll\s+Items\s+are\s+(AW\d{2})\s+version\b",
        ]
        for pat in patterns:
            m = re.search(pat, text, re.IGNORECASE)
            if m:
                return m.group(1).upper()
        return ""

    @staticmethod
    def _parse_reference(ref: str) -> tuple[str, str]:
        if not ref:
            return "", ""
        customer = ""
        cpo = ""
        bracket_tail = ""
        bracket = re.search(r"\[\s*([^\]]+?)\s*\]", ref)
        if bracket:
            customer = bracket.group(1).strip().upper()
            bracket_tail = ref[bracket.end():]
        else:
            lead = ref.strip().split()[0].upper() if ref.strip().split() else ""
            if lead == "UK":
                customer = "CASDON UK"
            elif lead in ("USA", "US"):
                customer = "CASDON USA"
            elif lead == "EMMA":
                customer = "CASDON EMMA"
            elif lead in ("EMEA", "EU"):
                customer = "CASDON EMEA"
            elif "AMAZON" in ref.upper():
                customer = "AMAZON"
        cpo_m = re.search(r"\b(P\d+)\b", ref, re.IGNORECASE)
        if cpo_m:
            cpo = cpo_m.group(1).upper()
        elif bracket_tail:
            # Bracketed references put the customer PO after "[CUSTOMER]", e.g.
            # "DAM [TARGET] 10001769117" / "DUK [AMAZON] 7CU324UE". Container
            # sizes (20FT/40FTHC) and short codes are not customer POs.
            candidates = []
            for token in re.findall(r"[A-Z0-9]{6,}", bracket_tail.upper()):
                if re.fullmatch(r"\d{1,3}FT(?:HC)?", token):
                    continue
                if sum(ch.isdigit() for ch in token) >= 4:
                    candidates.append(token)
            cpo = "/".join(dict.fromkeys(candidates))
        return customer, cpo

    @staticmethod
    def _parse_container_type(ref: str) -> str:
        text = re.sub(r"\s+", "", str(ref or "").upper())
        match = re.search(r"(20FT(?:HC)?|40FT(?:HC)?|45FT(?:HC)?|LCL|FCL)", text)
        return match.group(1) if match else ""

    def _parse_lines(self, text: str) -> list[CasdonLine]:
        parsed: list[CasdonLine] = []
        allow_continuation = False
        for raw in self._candidate_lines(text):
            line = self._parse_product_line(raw) or self._parse_charge_line(raw)
            if line:
                parsed.append(line)
                allow_continuation = not line.is_charge
            elif parsed and allow_continuation and self._is_description_continuation(raw):
                parsed[-1].description_en = f"{parsed[-1].description_en} {raw}".strip()
                allow_continuation = False
            elif re.match(r"^\d+\s+", raw):
                allow_continuation = False
        return parsed

    @staticmethod
    def _candidate_lines(text: str) -> Iterable[str]:
        for raw in text.splitlines():
            line = re.sub(r"\s+", " ", raw.replace("\xa0", " ")).strip()
            if not line:
                continue
            yield line

    def _parse_product_line(self, line: str) -> CasdonLine | None:
        m = self.PRODUCT_RE.match(line)
        if not m:
            return None
        return CasdonLine(
            line_no=m.group(1),
            item_code=m.group(2).strip(),
            description_en=m.group(3).strip(),
            commodity_code=m.group(4),
            qty=_int(m.group(5)),
            unit=m.group(6),
            unit_price=_num(m.group(7)),
            total_usd=_num(m.group(8)),
        )

    def _parse_charge_line(self, line: str) -> CasdonLine | None:
        m = self.CHARGE_RE.match(line)
        if not m:
            return None
        desc = m.group(2).strip()
        # Notes such as "This Order ... 0 0.00 0.00" are intentionally ignored.
        if _num(m.group(5)) <= 0 and _num(m.group(6)) <= 0:
            return None
        return CasdonLine(
            line_no=m.group(1),
            item_code="额外运费",
            description_en=desc,
            qty=_int(m.group(3)),
            unit=m.group(4),
            unit_price=_num(m.group(5)),
            total_usd=_num(m.group(6)),
            is_charge=True,
        )

    @staticmethod
    def _is_description_continuation(line: str) -> bool:
        if re.match(r"^\d+\s+", line):
            return False
        if re.match(r"^(TOTAL|Deliver To|This order|on behalf|Supplier:|Purchase Order)\b", line, re.I):
            return False
        if re.search(r"\b(?:NET|VAT|GROSS)\b", line):
            return False
        return bool(re.search(r"[A-Za-z]", line))


def validate(parsed: dict, filename: str = "") -> list[str]:
    warnings = []
    if not parsed.get("po_number"):
        warnings.append(f"{filename}: 未识别到合同号/PO号")
    missing_header = []
    if not parsed.get("po_date"):
        missing_header.append("Order Date")
    if not parsed.get("ship_date"):
        missing_header.append("Shipment Date")
    if not parsed.get("customer"):
        missing_header.append("Your Reference/客名")
    if missing_header:
        if str(filename).lower().endswith((".xlsx", ".xlsm")):
            warnings.append(
                f"{filename}: WPS转换Excel未保留抬头字段（{'、'.join(missing_header)}）；"
                "产品明细仍会识别，上传原PDF可补齐抬头"
            )
        else:
            warnings.append(f"{filename}: 未识别到{'、'.join(missing_header)}，请人工核对")

    # 自洽校验一：逐行 数量×单价 = 行金额（容差1分钱）
    lines = parsed.get("lines") or []
    for line in lines:
        if line.get("is_charge"):
            continue
        qty, price, total = line.get("qty") or 0, line.get("unit_price") or 0, line.get("total_usd") or 0
        if qty > 0 and price > 0:
            expect = round(qty * price, 2)
            if abs(expect - total) > 0.011:
                warnings.append(
                    f"{filename}: 行{line.get('line_no')} {line.get('item_code')} "
                    f"数量×单价({qty}×{price}={expect})与行金额({total})不符，请人工核对"
                )

    # 自洽校验二：所有行合计 = 单据 TOTAL NET AMOUNT
    raw = parsed.get("raw_text") or ""
    m = re.search(r"TOTAL\s+NET\s+AMOUNT\s+USD?\s*([\d,]+\.\d{2})", raw, re.IGNORECASE)
    if m and lines:
        doc_total = _num(m.group(1))
        line_sum = round(sum(l.get("total_usd") or 0 for l in lines), 2)
        if abs(doc_total - line_sum) > 0.011:
            warnings.append(
                f"{filename}: 行合计({line_sum})与单据TOTAL({doc_total})不符——"
                "可能有明细行未识别或金额识别错误，请人工核对"
            )
    return warnings
