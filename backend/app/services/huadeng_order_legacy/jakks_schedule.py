# -*- coding: utf-8 -*-
from __future__ import annotations

import csv
import io
import os
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable

import xlrd
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get("JAKKS_DATA_DIR", BASE_DIR / "华登-Jakks客AI"))
UPLOAD_DIR = Path(os.environ.get("JAKKS_UPLOAD_DIR", BASE_DIR / "uploads"))
WORKBOOK_SUFFIXES = {".xls", ".xlsx"}
REQUIRED_HEADER_FIELDS = {"customer_po", "contract_no", "sku", "quantity"}
MIN_RECOGNIZED_HEADERS = 8


@dataclass(frozen=True)
class WorkbookCell:
    value: Any
    number_format: str = ""
    is_date: bool = False

FIELD_MAP = {
    "来单日期": "order_date",
    "最终改单日期": "final_change_date",
    "合同联系人": "contact",
    "客户PO": "customer_po",
    "合同号": "contract_no",
    "客名": "customer",
    "版本": "version",
    "货号": "sku",
    "产品名称": "product",
    "数量": "quantity",
    "内装箱": "inner_carton",
    "外装箱": "outer_carton",
    "总箱数": "cartons",
    "特别备注": "special_note",
    "箱唛": "carton_mark",
    "客供标签": "customer_label",
    "对应MO No.": "mo_no",
    "来料复期": "material_reply_date",
    "胶件改动": "plastic_change",
    "纸箱复期": "carton_reply_date",
    "注料复期": "injection_reply_date",
    "上拉期": "start_date",
    "完成期": "finish_date",
    "日期码": "date_code",
    "验货期": "inspection_date",
    "客验期": "customer_inspection_date",
    "PO走货期": "ship_date",
    "预计订仓时间": "booking_estimate_date",
    "收货地点（装箱港）": "receiving_port",
    "单价（港币）": "unit_hkd",
    "单价(美金)": "unit_usd",
    "总货价（港币）": "total_hkd",
    "总货价（美金）": "total_usd",
    "报价装配工HK$": "assembly_labor_hkd",
    "品牌名": "brand",
    "预追客供物料期": "material_follow_date",
    "入系统时间": "system_date",
    "外箱 长CM": "outer_length_cm",
    "外箱 宽CM": "outer_width_cm",
    "外箱 高CM": "outer_height_cm",
    "CBM /M3": "cbm",
    "毛重KG": "gross_weight_kg",
    "净重KG": "net_weight_kg",
    "总CBM /M3": "total_cbm",
    "走货国家": "country",
    "预计SO订舱时间": "so_booking_date",
    "确认号SH": "confirmation_no",
    "走货方式": "shipping_method",
    "发票号码": "invoice_no",
    "发票日期": "invoice_date",
    "彩盒条码货号": "barcode_item",
    "JAKKS/TOLLY TOY/DISGUISE": "jakks_category",
    "柜号/车": "container_no",
    "生产地点": "production_site",
    "更改为题点": "change_point",
}

DATE_FIELDS = {
    "order_date",
    "final_change_date",
    "material_reply_date",
    "carton_reply_date",
    "injection_reply_date",
    "start_date",
    "finish_date",
    "inspection_date",
    "customer_inspection_date",
    "ship_date",
    "booking_estimate_date",
    "material_follow_date",
    "system_date",
    "so_booking_date",
    "invoice_date",
}

EXPORT_FIELDS = [
    ("状态", "status"),
    ("风险", "risk_text"),
    ("来单日期", "order_date"),
    ("联系人", "contact"),
    ("客户PO", "customer_po"),
    ("合同号", "contract_no"),
    ("客户", "customer"),
    ("货号", "sku"),
    ("产品名称", "product"),
    ("数量", "quantity"),
    ("总箱数", "cartons"),
    ("品牌", "brand"),
    ("国家", "country"),
    ("纸箱复期", "carton_reply_date"),
    ("上拉期", "start_date"),
    ("完成期", "finish_date"),
    ("验货期", "inspection_date"),
    ("客验期", "customer_inspection_date"),
    ("PO走货期", "ship_date"),
    ("预计SO订舱时间", "so_booking_date"),
    ("走货方式", "shipping_method"),
    ("确认号SH", "confirmation_no"),
    ("特别备注", "special_note"),
]


def find_default_workbook() -> Path:
    """Return the newest workbook from uploads first, then bundled data."""
    for folder in (UPLOAD_DIR, DATA_DIR, BASE_DIR):
        if not folder.exists():
            continue
        candidates = [
            path
            for path in folder.iterdir()
            if path.is_file() and path.suffix.lower() in WORKBOOK_SUFFIXES
        ]
        if candidates:
            return max(candidates, key=lambda path: path.stat().st_mtime)
    raise FileNotFoundError("未找到 .xls 或 .xlsx 当前有效排期；首次使用或业务换版时上传一次即可")


def list_source_files() -> dict[str, list[dict[str, Any]]]:
    files: dict[str, list[dict[str, Any]]] = {"workbooks": [], "pdfs": []}
    for folder in (UPLOAD_DIR, DATA_DIR):
        if not folder.exists():
            continue
        for path in sorted(folder.iterdir(), key=lambda item: item.stat().st_mtime, reverse=True):
            if path.suffix.lower() in WORKBOOK_SUFFIXES:
                files["workbooks"].append(_file_info(path))
            elif path.suffix.lower() == ".pdf":
                files["pdfs"].append(_file_info(path))
    return files


def load_dataset(path: str | Path | None = None, today: date | None = None) -> dict[str, Any]:
    workbook_path = Path(path) if path else find_default_workbook()
    report_date = today or date.today()
    records = parse_workbook(workbook_path, report_date)
    summary = summarize(records)
    return {
        "source_file": _file_info(workbook_path),
        "source_files": list_source_files(),
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "today": report_date.isoformat(),
        "summary": summary,
        "options": build_options(records),
        "schedule": build_schedule(records),
        "records": records,
    }


def parse_workbook(path: str | Path, today: date | None = None) -> list[dict[str, Any]]:
    workbook_path = Path(path)
    if workbook_path.suffix.lower() not in WORKBOOK_SUFFIXES:
        raise ValueError("仅支持 .xls 或 .xlsx 排期表")

    rows, datemode, epoch = _read_workbook(workbook_path)
    header_index = _find_header_row(rows)
    header_row = [_display_cell(cell) for cell in rows[header_index]]
    column_map = {
        index: FIELD_MAP[header]
        for index, header in enumerate(header_row)
        if header in FIELD_MAP
    }
    default_year = _guess_year(workbook_path)
    report_date = today or date.today()

    records: list[dict[str, Any]] = []
    for row_index in range(header_index + 1, len(rows)):
        record: dict[str, Any] = {
            "id": f"r{row_index + 1}",
            "row_no": row_index + 1,
            "_dates": {},
            "_raw": {},
        }
        for col_index, field in column_map.items():
            cell = rows[row_index][col_index] if col_index < len(rows[row_index]) else WorkbookCell("")
            value = cell.value
            display_value = _display_cell(cell)
            raw_value = display_value
            if field in DATE_FIELDS:
                raw_value, iso_value = _date_cell_value(cell, datemode, epoch, default_year)
                record["_dates"][field] = iso_value
            record[field] = raw_value
            record["_raw"][field] = display_value

        if _should_skip(record):
            continue

        record["quantity_num"] = _number_value(record.get("quantity"))
        record["cartons_num"] = _number_value(record.get("cartons"))
        record["total_usd_num"] = _number_value(record.get("total_usd"))
        record["total_hkd_num"] = _number_value(record.get("total_hkd"))
        record["week_key"], record["week_label"] = _week_label(record)
        status = assess_record(record, report_date)
        record.update(status)
        records.append(record)

    if not records:
        raise ValueError("排期表未包含可识别的订单数据")
    return records


def _read_workbook(
    path: Path,
) -> tuple[list[list[WorkbookCell]], int | None, Any | None]:
    if path.suffix.lower() == ".xlsx":
        book = load_workbook(path, read_only=False, data_only=True)
        sheet = _pick_xlsx_sheet(book)
        rows = [
            [
                WorkbookCell(cell.value, str(cell.number_format or ""), bool(cell.is_date))
                for cell in row
            ]
            for row in sheet.iter_rows()
        ]
        return rows, None, book.epoch

    book = xlrd.open_workbook(str(path), formatting_info=False)
    sheet = _pick_sheet(book)
    rows = [
        [WorkbookCell(sheet.cell_value(row_index, col_index)) for col_index in range(sheet.ncols)]
        for row_index in range(sheet.nrows)
    ]
    return rows, book.datemode, None


def _find_header_row(rows: list[list[WorkbookCell]]) -> int:
    best_index = -1
    best_fields: set[str] = set()
    for index, row in enumerate(rows[:20]):
        fields = {
            FIELD_MAP[label]
            for label in (_display_cell(cell) for cell in row)
            if label in FIELD_MAP
        }
        if len(fields) > len(best_fields):
            best_index = index
            best_fields = fields

    if (
        best_index < 0
        or len(best_fields) < MIN_RECOGNIZED_HEADERS
        or not REQUIRED_HEADER_FIELDS.issubset(best_fields)
    ):
        missing = sorted(REQUIRED_HEADER_FIELDS - best_fields)
        detail = f"，缺少核心字段：{', '.join(missing)}" if missing else ""
        raise ValueError(
            f"不是有效的 Jakks 排期表：仅识别到 {len(best_fields)} 个排期字段{detail}"
        )
    return best_index


def assess_record(record: dict[str, Any], today: date) -> dict[str, Any]:
    risks: list[dict[str, str]] = []
    ship_date = _iso_date(record, "ship_date")
    finish_date = _iso_date(record, "finish_date")
    inspection_date = _iso_date(record, "inspection_date")
    customer_inspection_date = _iso_date(record, "customer_inspection_date")
    booking_date = _iso_date(record, "so_booking_date") or _iso_date(record, "booking_estimate_date")
    has_shipping_marker = bool(record.get("invoice_no") or record.get("container_no"))

    if not ship_date:
        risks.append({"code": "missing_ship", "label": "缺走货期", "level": "high"})
    elif ship_date < today and not has_shipping_marker:
        days = (today - ship_date).days
        risks.append({"code": "overdue_ship", "label": f"走货逾期{days}天", "level": "high"})
    elif ship_date <= today + timedelta(days=7):
        days = max((ship_date - today).days, 0)
        risks.append({"code": "ship_soon", "label": f"{days}天内走货", "level": "medium"})

    if ship_date and finish_date and finish_date > ship_date:
        risks.append({"code": "finish_after_ship", "label": "完成期晚于走货", "level": "high"})
    if ship_date and booking_date and booking_date > ship_date:
        risks.append({"code": "booking_after_ship", "label": "订舱晚于走货", "level": "high"})
    if ship_date and inspection_date and inspection_date > ship_date:
        risks.append({"code": "inspection_after_ship", "label": "验货晚于走货", "level": "medium"})
    if ship_date and customer_inspection_date and customer_inspection_date > ship_date:
        risks.append({"code": "customer_inspection_after_ship", "label": "客验晚于走货", "level": "medium"})

    if ship_date and ship_date >= today:
        if not finish_date:
            risks.append({"code": "missing_finish", "label": "缺完成期", "level": "medium"})
        if not (inspection_date or customer_inspection_date):
            risks.append({"code": "missing_inspection", "label": "缺验货/客验", "level": "medium"})
        if not booking_date:
            risks.append({"code": "missing_booking", "label": "缺订舱期", "level": "medium"})

    material_text = " ".join(
        str(record.get(field, ""))
        for field in ("special_note", "carton_mark", "customer_label")
    )
    if "欠" in material_text:
        risks.append({"code": "material_pending", "label": "物料/标签欠", "level": "medium"})

    if _number_value(record.get("quantity")) is None:
        risks.append({"code": "missing_quantity", "label": "缺数量", "level": "low"})

    level_rank = {"high": 3, "medium": 2, "low": 1}
    max_level = max((level_rank[risk["level"]] for risk in risks), default=0)
    if max_level >= 3:
        status = "严重"
        status_class = "danger"
    elif max_level == 2:
        status = "关注"
        status_class = "warn"
    elif max_level == 1:
        status = "提示"
        status_class = "info"
    else:
        status = "正常"
        status_class = "ok"

    return {
        "status": status,
        "status_class": status_class,
        "risks": risks,
        "risk_text": "、".join(risk["label"] for risk in risks) or "无",
    }


def summarize(records: list[dict[str, Any]]) -> dict[str, Any]:
    risk_counter = Counter()
    status_counter = Counter(record["status"] for record in records)
    for record in records:
        for risk in record["risks"]:
            risk_counter[risk["code"]] += 1

    ship_dates = [_iso_date(record, "ship_date") for record in records]
    ship_dates = [ship_date for ship_date in ship_dates if ship_date]
    next_ship = min(ship_dates).isoformat() if ship_dates else ""
    last_ship = max(ship_dates).isoformat() if ship_dates else ""

    return {
        "total_orders": len(records),
        "total_quantity": _sum_field(records, "quantity_num"),
        "total_cartons": _sum_field(records, "cartons_num"),
        "total_usd": round(_sum_field(records, "total_usd_num"), 2),
        "total_hkd": round(_sum_field(records, "total_hkd_num"), 2),
        "countries": len({record.get("country") for record in records if record.get("country")}),
        "brands": len({record.get("brand") for record in records if record.get("brand")}),
        "status": dict(status_counter),
        "risks": dict(risk_counter),
        "next_ship_date": next_ship,
        "last_ship_date": last_ship,
    }


def build_options(records: list[dict[str, Any]]) -> dict[str, list[str]]:
    return {
        "contacts": _sorted_values(record.get("contact") for record in records),
        "brands": _sorted_values(record.get("brand") for record in records),
        "countries": _sorted_values(record.get("country") for record in records),
        "methods": _sorted_values(record.get("shipping_method") for record in records),
        "statuses": ["严重", "关注", "提示", "正常"],
        "risks": _sorted_values(risk["label"] for record in records for risk in record["risks"]),
    }


def build_schedule(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    buckets: dict[str, dict[str, Any]] = {}
    for record in records:
        key = record.get("week_key") or "未排期"
        label = record.get("week_label") or "未排期"
        if key not in buckets:
            buckets[key] = {
                "key": key,
                "label": label,
                "count": 0,
                "quantity": 0.0,
                "cartons": 0.0,
                "danger": 0,
                "warn": 0,
                "ok": 0,
                "countries": Counter(),
                "brands": Counter(),
            }
        bucket = buckets[key]
        bucket["count"] += 1
        bucket["quantity"] += record.get("quantity_num") or 0
        bucket["cartons"] += record.get("cartons_num") or 0
        bucket[record.get("status_class", "ok")] += 1
        if record.get("country"):
            bucket["countries"][record["country"]] += 1
        if record.get("brand"):
            bucket["brands"][record["brand"]] += 1

    schedule: list[dict[str, Any]] = []
    for key in sorted(buckets):
        bucket = buckets[key]
        bucket["quantity"] = round(bucket["quantity"], 2)
        bucket["cartons"] = round(bucket["cartons"], 2)
        bucket["top_country"] = _top_label(bucket.pop("countries"))
        bucket["top_brand"] = _top_label(bucket.pop("brands"))
        schedule.append(bucket)
    return schedule


def apply_filters(records: Iterable[dict[str, Any]], filters: dict[str, str]) -> list[dict[str, Any]]:
    result = list(records)
    search = filters.get("search", "").strip().lower()
    if search:
        fields = ("customer_po", "contract_no", "customer", "sku", "product", "country", "brand", "confirmation_no")
        result = [
            record
            for record in result
            if any(search in str(record.get(field, "")).lower() for field in fields)
        ]
    for key, field in (
        ("contact", "contact"),
        ("brand", "brand"),
        ("country", "country"),
        ("method", "shipping_method"),
        ("status", "status"),
    ):
        value = filters.get(key, "").strip()
        if value:
            result = [record for record in result if str(record.get(field, "")) == value]
    risk = filters.get("risk", "").strip()
    if risk:
        result = [
            record
            for record in result
            if any(item["label"] == risk or item["code"] == risk for item in record["risks"])
        ]
    return result


def records_to_csv(records: Iterable[dict[str, Any]]) -> str:
    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow([label for label, _ in EXPORT_FIELDS])
    for record in records:
        writer.writerow([record.get(field, "") for _, field in EXPORT_FIELDS])
    return output.getvalue()


def _pick_sheet(book: xlrd.Book) -> xlrd.sheet.Sheet:
    for sheet in book.sheets():
        if "排货" in sheet.name or "schedule" in sheet.name.lower():
            return sheet
    return book.sheet_by_index(0)


def _pick_xlsx_sheet(book: Any) -> Any:
    for sheet in book.worksheets:
        if "排货" in sheet.title or "schedule" in sheet.title.lower():
            return sheet
    return book.worksheets[0]


def _should_skip(record: dict[str, Any]) -> bool:
    product = str(record.get("product", "")).strip()
    if product == "合计":
        return True
    meaningful = [
        record.get("customer_po"),
        record.get("contract_no"),
        record.get("customer"),
        record.get("sku"),
        record.get("quantity"),
    ]
    if not any(value not in (None, "") and str(value).strip() for value in meaningful):
        return True
    if not record.get("sku") and product.startswith("以下"):
        return True
    return False


def _date_cell_value(
    cell: WorkbookCell,
    datemode: int | None,
    epoch: Any | None,
    default_year: int,
) -> tuple[str, str]:
    value = cell.value
    if value in ("", None):
        return "", ""
    if isinstance(value, datetime):
        parsed = value.date()
        return parsed.isoformat(), parsed.isoformat()
    if isinstance(value, date):
        return value.isoformat(), value.isoformat()
    if isinstance(value, (int, float)) and value:
        try:
            if datemode is not None:
                parsed = xlrd.xldate_as_datetime(float(value), datemode).date()
            elif epoch is not None:
                converted = from_excel(float(value), epoch)
                parsed = converted.date() if isinstance(converted, datetime) else converted
            else:
                raise ValueError("缺少 Excel 日期系统")
            return parsed.isoformat(), parsed.isoformat()
        except (TypeError, ValueError, OverflowError):
            return _display_cell(cell), ""

    text = str(value).strip()
    parsed = _parse_text_date(text, default_year)
    if parsed:
        return parsed.isoformat(), parsed.isoformat()
    return text, ""


def _display_cell(cell: WorkbookCell) -> str:
    value = cell.value
    if value in ("", None):
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        cleaned_format = re.sub(r"_.|\\.", "", cell.number_format.split(";")[0]).strip()
        if value >= 0 and float(value).is_integer() and re.fullmatch(r"0+", cleaned_format):
            return str(int(value)).zfill(len(cleaned_format))
    return _display_value(value)


def _parse_text_date(text: str, default_year: int) -> date | None:
    full = re.search(r"(20\d{2})[./-](\d{1,2})[./-](\d{1,2})", text)
    if full:
        return _safe_date(int(full.group(1)), int(full.group(2)), int(full.group(3)))

    compact = re.search(r"(?<!\d)(\d{1,2})[./-](\d{1,2})(?!\d)", text)
    if compact:
        return _safe_date(default_year, int(compact.group(1)), int(compact.group(2)))
    return None


def _safe_date(year: int, month: int, day: int) -> date | None:
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _display_value(value: Any) -> str:
    if value in ("", None):
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.6f}".rstrip("0").rstrip(".")
    return str(value).strip()


def _number_value(value: Any) -> float | None:
    if value in ("", None):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _iso_date(record: dict[str, Any], field: str) -> date | None:
    value = record.get("_dates", {}).get(field)
    if not value:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError:
        return None


def _week_label(record: dict[str, Any]) -> tuple[str, str]:
    target = (
        _iso_date(record, "ship_date")
        or _iso_date(record, "customer_inspection_date")
        or _iso_date(record, "finish_date")
        or _iso_date(record, "order_date")
    )
    if not target:
        return "未排期", "未排期"
    start = target - timedelta(days=target.weekday())
    end = start + timedelta(days=6)
    iso_year, iso_week, _ = target.isocalendar()
    key = f"{iso_year}-W{iso_week:02d}"
    label = f"{key} {start:%m/%d}-{end:%m/%d}"
    return key, label


def _guess_year(path: Path) -> int:
    match = re.search(r"(20\d{2})", path.name)
    if match:
        return int(match.group(1))
    return date.today().year


def _sum_field(records: Iterable[dict[str, Any]], field: str) -> float:
    return round(sum(record.get(field) or 0 for record in records), 2)


def _sorted_values(values: Iterable[Any]) -> list[str]:
    return sorted({str(value).strip() for value in values if str(value).strip()})


def _top_label(counter: Counter) -> str:
    if not counter:
        return ""
    label, count = counter.most_common(1)[0]
    return f"{label} ({count})"


def _file_info(path: Path) -> dict[str, Any]:
    stat = path.stat()
    return {
        "name": path.name,
        "path": str(path),
        "size": stat.st_size,
        "modified": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
    }
