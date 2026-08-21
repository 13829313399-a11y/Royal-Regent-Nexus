# -*- coding: utf-8 -*-
"""Casdon single schedule matching, updating, and new-order export."""
from __future__ import annotations

import os
import re
import shutil
from collections import Counter, defaultdict
from copy import copy
from datetime import datetime, timedelta
from typing import Any

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter


MASTER_SHEET = "Casdon 排货表-总"
MASTER_SHEET_CANDIDATES = (
    "2026年未验货订单",
    MASTER_SHEET,
)
WAREHOUSE_TEMPLATE_SHEET = "casdon仓单PO模版"

BLUE_FILL = PatternFill(start_color="00B0F0", end_color="00B0F0", fill_type="solid")

COL = {
    "po_date": 2,
    "change_note": 3,
    "contact": 4,
    "customer_po": 5,
    "contract": 6,
    "customer": 7,
    "version": 8,
    "item": 9,
    "cn_name": 10,
    "qty": 11,
    "inner": 12,
    "outer": 13,
    "total_box": 14,
    "special_note": 15,
    "carton_mark": 16,
    "customer_label": 17,
    "customer_material": 18,
    "material_reply": 19,
    "plastic_reply": 20,
    "carton_reply": 21,
    "injection_reply": 22,
    "line_start": 23,
    "complete_date": 24,
    "date_code": 25,
    "workshop_recheck": 26,
    "inspection_date": 27,
    "customer_inspection_date": 28,
    "ship_date": 29,
    "booking_date": 30,
    "delivery_place": 31,
    "unit_hkd": 32,
    "unit_usd": 33,
    "total_hkd": 34,
    "total_usd": 35,
    "assembly_quote_hkd": 36,
    "english_name": 37,
    "pre_customer_material_date": 38,
    "certificate": 39,
    "carton_l": 40,
    "carton_w": 41,
    "carton_h": 42,
    "carton_cbm": 43,
    "total_cbm": 44,
    "gross_weight": 45,
    "net_weight": 46,
    "ship_country": 47,
    "so_no": 48,
    "container_type": 49,
    "warehouse_date": 50,
    "truck_detail": 51,
    "brand": 52,
}

EXPORT_END_COL = 79  # CA
CORE_END_COL = 30    # AD
HKD_RATE = 7.75


def excel_serial(dt: datetime | None) -> int | None:
    if not dt:
        return None
    return (dt - datetime(1899, 12, 30)).days


def parse_date(value: Any) -> datetime | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value
    text = str(value)[:10]
    try:
        return datetime.strptime(text, "%Y-%m-%d")
    except ValueError:
        return None


def item_base(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper().replace("（", "(")
    if text == "额外运费":
        return "额外运费"
    text = re.sub(r"\s+", "", text)
    m = re.match(r"([A-Z]*\d+[A-Z0-9./-]*)", text)
    if not m:
        return ""
    base = m.group(1)
    # 11050.TAR001 should still match a PDF item code 11050.
    m2 = re.match(r"(\d+)", base)
    return m2.group(1) if m2 else base


def normalize_po(value: Any) -> str:
    text = str(value or "").strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text


def _display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = str(value)
    return text[:-2] if text.endswith(".0") else text


def _value_equal(old: Any, new: Any) -> bool:
    if old is None and new in (None, ""):
        return True
    if isinstance(old, datetime) and isinstance(new, int):
        return excel_serial(old) == new
    try:
        if old not in (None, "") and new not in (None, ""):
            # 1分钱内视为相同：人工常填全精度乘积(10794.924)，PO单据是两位小数(10794.92)，
            # 这类尾数差异不算改动；日期序列/数量都是整数，差1仍会正常识别为变化。
            return abs(float(old) - float(new)) < 0.011
    except (TypeError, ValueError):
        pass
    return _display(old).strip() == _display(new).strip()


class MasterIndex:
    def __init__(self):
        # (合同号, 货号base) -> [(行号, 该行数量), ...]；同一 PO 可能分批出货，
        # 总表允许多行同键，匹配时按数量配对。
        self.index: dict[tuple[str, str], list[tuple[int, Any]]] = {}
        self.catalog_by_item: dict[str, dict] = {}
        self.catalog_by_customer_item: dict[tuple[str, str], dict] = {}
        self.metadata_by_contract_item: dict[tuple[str, str], dict] = {}
        self.max_row = 1
        self.headers: dict[int, str] = {}
        self.inheritance_conflicts: list[dict[str, Any]] = []


def _detect_header_row(ws) -> int:
    """Find the actual Casdon order-table header instead of trusting a legacy tab name."""
    required = ("合同号", "货号", "产品名称", "数量")
    for row_no in range(1, min(ws.max_row or 1, 12) + 1):
        values = {
            str(ws.cell(row_no, col).value or "").replace("\n", "").replace(" ", "")
            for col in range(1, min(ws.max_column or EXPORT_END_COL, EXPORT_END_COL) + 1)
        }
        if all(any(token in value for value in values) for token in required):
            return row_no
    return 0


def _select_master_sheet(wb):
    """Select the live order sheet and never fall back to RFID/notes/history tabs."""
    for name in MASTER_SHEET_CANDIDATES:
        if name in wb.sheetnames:
            ws = wb[name]
            header_row = _detect_header_row(ws)
            if header_row:
                return ws, header_row

    excluded = ("取消", "历史", "RFID", "注意", "说明")
    for ws in wb.worksheets:
        if any(token.lower() in ws.title.lower() for token in excluded):
            continue
        header_row = _detect_header_row(ws)
        if header_row:
            return ws, header_row

    raise ValueError(
        "Casdon 排期中未找到含“合同号、货号、产品名称、数量”的有效订单表，"
        f"现有工作表：{'、'.join(wb.sheetnames)}"
    )


def build_index(filepath: str) -> MasterIndex:
    idx = MasterIndex()
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    try:
        ws, header_row = _select_master_sheet(wb)
        idx.max_row = ws.max_row
        idx.headers = {
            c: str(ws.cell(header_row, c).value or "")
            for c in range(1, min(EXPORT_END_COL, ws.max_column) + 1)
        }
        counters = defaultdict(lambda: defaultdict(Counter))
        samples = defaultdict(dict)
        customer_counters = defaultdict(lambda: defaultdict(Counter))
        customer_samples = defaultdict(dict)

        for r, row in enumerate(
            ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row, max_col=EXPORT_END_COL),
            start=header_row + 1,
        ):
            po = normalize_po(row[COL["contract"] - 1].value)
            item_value = row[COL["item"] - 1].value
            base = item_base(item_value)
            customer = str(row[COL["customer"] - 1].value or "").strip().upper()
            if po and base:
                idx.index.setdefault((po, base), []).append((r, row[COL["qty"] - 1].value))
                exact = idx.metadata_by_contract_item.setdefault((po, base), {})
                for key, col in _CATALOG_COLS.items():
                    value = row[col - 1].value
                    if value not in (None, "") and key not in exact:
                        exact[key] = value
                for key, col in (
                    ("customer_po", COL["customer_po"]),
                    ("ship_date", COL["ship_date"]),
                ):
                    value = row[col - 1].value
                    if value not in (None, "") and key not in exact:
                        exact[key] = value
            if not base:
                continue
            for key, col in _CATALOG_COLS.items():
                val = row[col - 1].value
                if val not in (None, ""):
                    counters[base][key].update([val])
                    if key not in samples[base]:
                        samples[base][key] = val
                    if customer:
                        customer_counters[(customer, base)][key].update([val])
                        if key not in customer_samples[(customer, base)]:
                            customer_samples[(customer, base)][key] = val

        for base in counters:
            idx.catalog_by_item[base] = _counter_to_catalog(
                counters[base], samples[base], idx.inheritance_conflicts, "全部客户", base
            )
        for key in customer_counters:
            idx.catalog_by_customer_item[key] = _counter_to_catalog(
                customer_counters[key], customer_samples[key], idx.inheritance_conflicts, key[0], key[1]
            )
        return idx
    finally:
        wb.close()


_CATALOG_COLS = {
    "contact": COL["contact"],
    "customer": COL["customer"],
    "version": COL["version"],
    "item_value": COL["item"],
    "cn_name": COL["cn_name"],
    "english_name": COL["english_name"],
    "inner": COL["inner"],
    "outer": COL["outer"],
    "special_note": COL["special_note"],
    "carton_mark": COL["carton_mark"],
    "customer_label": COL["customer_label"],
    "customer_material": COL["customer_material"],
    "carton_reply": COL["carton_reply"],
    "delivery_place": COL["delivery_place"],
    "carton_l": COL["carton_l"],
    "carton_w": COL["carton_w"],
    "carton_h": COL["carton_h"],
    "carton_cbm": COL["carton_cbm"],
    "gross_weight": COL["gross_weight"],
    "net_weight": COL["net_weight"],
    "ship_country": COL["ship_country"],
    "brand": COL["brand"],
}


_PRODUCT_NAME_FIELDS = {"cn_name", "english_name"}


def _name_key(value: Any) -> str:
    return re.sub(r"[\W_]+", "", str(value or "").strip().casefold())


def _counter_to_catalog(
    counter_map: dict[str, Counter],
    samples: dict,
    conflicts: list[dict[str, Any]],
    scope: str,
    item: str,
) -> dict:
    result = {}
    for key, counter in counter_map.items():
        if key in _PRODUCT_NAME_FIELDS:
            values = {}
            for value in counter:
                normalized = _name_key(value)
                if normalized:
                    values.setdefault(normalized, value)
            if len(values) == 1:
                result[key] = next(iter(values.values()))
            elif len(values) > 1:
                conflicts.append({
                    "scope": scope,
                    "item": item,
                    "field": key,
                    "values": [str(value) for value in list(values.values())[:8]],
                })
            continue
        try:
            result[key] = counter.most_common(1)[0][0]
        except TypeError:
            result[key] = samples.get(key)
    return result


def _customer_family(value: Any) -> str:
    text = re.sub(r"[^A-Z0-9]+", " ", str(value or "").upper()).strip()
    for token in ("TARGET", "WALMART", "AMAZON", "COSTCO", "WHB", "KIDOODLE"):
        if token in text:
            return token
    if "CASDON" in text:
        for region in ("UK", "USA", "EMMA", "EMEA"):
            if region in text:
                return f"CASDON {region}"
        return "CASDON"
    return text


def _catalog_for(idx: MasterIndex, customer: str, base: str, po: str = "") -> dict:
    key = (str(customer or "").strip().upper(), base)
    # 品名可按货号作为产品主数据继承；联系人、版本、包装、标签等客户资料
    # 只能从同一客名+货号继承，防止同货号跨客户串用。
    global_catalog = idx.catalog_by_item.get(base, {})
    cat = {
        name: global_catalog[name]
        for name in _PRODUCT_NAME_FIELDS
        if global_catalog.get(name) not in (None, "")
    }
    exact = idx.metadata_by_contract_item.get((normalize_po(po), base), {})
    if exact:
        cat.update(exact)
        return cat

    cat.update(idx.catalog_by_customer_item.get(key, {}))
    if customer and len(cat) <= len(_PRODUCT_NAME_FIELDS):
        family = _customer_family(customer)
        matches = [
            catalog
            for (stored_customer, stored_base), catalog in idx.catalog_by_customer_item.items()
            if stored_base == base and _customer_family(stored_customer) == family
        ]
        if len(matches) == 1:
            cat.update(matches[0])
    return cat


def _previous_weekday(value: datetime | None) -> datetime | None:
    """Move Saturday/Sunday to the preceding Friday."""
    if not value:
        return None
    while value.weekday() >= 5:
        value -= timedelta(days=1)
    return value


def _date_code(item: Any, inspection_date: datetime | None) -> str:
    """Casdon: item first 3 + ROY + YYWW, using two weeks before inspection."""
    base = item_base(item)
    if not base or not inspection_date:
        return ""
    production_week = inspection_date - timedelta(days=14)
    iso = production_week.isocalendar()
    return f"{base[:3]}ROY{iso.year % 100:02d}{iso.week:02d}"


def _casdon_packaging(customer: str) -> tuple[str, str]:
    normalized = re.sub(r"\s+", " ", str(customer or "").strip().upper())
    if normalized in {"CASDON UK", "CASDON USA", "CASDON EMMA"}:
        return "标准唛", "无"
    return "欠箱唛", "欠利宝"


def _compose_row(order: dict, line: dict, idx: MasterIndex) -> dict[int, Any]:
    po = normalize_po(order.get("po_number"))
    customer = order.get("customer") or ""
    base = "额外运费" if line.get("is_charge") else item_base(line.get("item_code"))
    cat = _catalog_for(idx, customer, base, po)
    customer = cat.get("customer") or customer
    # 蓝色批注规定“来单日期”取正式 PO 邮件收件日，不得拿 PO Order Date 冒充。
    po_date = parse_date(order.get("email_received_date"))
    ship_date = parse_date(order.get("ship_date") or cat.get("ship_date"))
    insp_date = _previous_weekday(ship_date - timedelta(days=7)) if ship_date else None

    version = order.get("version") or cat.get("version") or ""
    item_value = "额外运费" if line.get("is_charge") else (cat.get("item_value") or _format_item(line.get("item_code"), version))
    cn_name = "额外运费" if line.get("is_charge") else cat.get("cn_name", "")
    qty = "" if line.get("is_charge") else line.get("qty", "")
    carton_mark, customer_label = _casdon_packaging(customer)
    unit_usd = line.get("unit_price", "") or ""
    total_usd = line.get("total_usd", "") or ""
    unit_hkd = round(float(unit_usd) * HKD_RATE, 6) if unit_usd not in (None, "") else ""
    total_hkd = round(float(total_usd) * HKD_RATE, 2) if total_usd not in (None, "") else ""

    values = {
        COL["po_date"]: excel_serial(po_date),
        COL["contact"]: cat.get("contact", ""),
        COL["customer_po"]: order.get("customer_po_header") or cat.get("customer_po", ""),
        COL["contract"]: po,
        COL["customer"]: customer or cat.get("customer", ""),
        COL["version"]: "" if line.get("is_charge") else version,
        COL["item"]: item_value,
        COL["cn_name"]: cn_name,
        COL["qty"]: qty,
        # 内箱没有明确继承规则，继续留空；外箱按反馈从同客户、同货号
        # 的排期主数据继承，总箱由导出层按“数量 / 外箱”写公式。
        COL["inner"]: "",
        COL["outer"]: "" if line.get("is_charge") else cat.get("outer", ""),
        COL["special_note"]: "",
        COL["carton_mark"]: "" if line.get("is_charge") else carton_mark,
        COL["customer_label"]: "" if line.get("is_charge") else customer_label,
        COL["customer_material"]: "",
        COL["carton_reply"]: "",
        COL["date_code"]: "" if line.get("is_charge") else _date_code(line.get("item_code"), insp_date),
        COL["inspection_date"]: excel_serial(insp_date),
        COL["ship_date"]: excel_serial(ship_date),
        COL["delivery_place"]: "",
        COL["unit_hkd"]: unit_hkd,
        COL["unit_usd"]: unit_usd,
        COL["total_hkd"]: total_hkd,
        COL["total_usd"]: total_usd,
        COL["english_name"]: (
            "" if line.get("is_charge")
            else line.get("description_en", "") or cat.get("english_name", "")
        ),
        COL["carton_l"]: "",
        COL["carton_w"]: "",
        COL["carton_h"]: "",
        COL["carton_cbm"]: "",
        COL["gross_weight"]: "",
        COL["net_weight"]: "",
        COL["ship_country"]: "",
        COL["container_type"]: order.get("container_type", ""),
        COL["brand"]: "",
    }
    return values


def _format_item(item_code: Any, version: str) -> str:
    code = str(item_code or "").strip()
    if not code:
        return ""
    return f"{code}({version})" if version else code


def write_orders(master_path: str, orders: list[dict], export_dir: str) -> dict:
    if not os.path.exists(master_path):
        return {"ok": False, "msg": f"排期文件不存在: {master_path}"}

    idx = build_index(master_path)
    new_rows = []
    new_details = []
    database_rows = []
    warnings = []
    inherited_rows = 0
    for order in orders:
        if not normalize_po(order.get("po_number")):
            warnings.append(f"{order.get('filename', '')}: 未识别合同号，请在导出表中人工复核")
        if not order.get("email_received_date"):
            warnings.append(
                f"{order.get('filename', '')}: “来单日期”应为正式 PO 邮件收件日，PO 内没有该字段，当前留空"
            )
        if order.get("ship_date"):
            warnings.append(
                f"{order.get('filename', '')}: 验货期已按走货期前7天并避开周末；法定节假日仍需业务复核"
            )
        for line in order.get("lines", []):
            values = _compose_row(order, line, idx)
            new_rows.append(values)
            new_details.append(values.get(COL["item"], line.get("item_code", "")))
            database_rows.append({
                "order_date": order.get("email_received_date") or order.get("po_date", ""),
                "po_no": order.get("your_reference") or order.get("customer_po_header", ""),
                "contract_no": order.get("po_number", ""),
                "client_name": order.get("customer", ""),
                "item_no": values.get(COL["item"], line.get("item_code", "")),
                "product_name": values.get(COL["cn_name"], "") or line.get("description_en", ""),
                "quantity": line.get("qty"),
                "unit": line.get("unit", ""),
                "inspection_date": values.get(COL["inspection_date"], ""),
                "ship_date": order.get("ship_date", ""),
            })
            if not line.get("is_charge") and not values.get(COL["cn_name"]):
                warnings.append(
                    f"{order.get('filename', '')} / {line.get('item_code', '')}: "
                    "当前排期未找到唯一中文品名，已留空"
                )
            if not line.get("is_charge") and (
                values.get(COL["cn_name"]) or values.get(COL["english_name"])
            ):
                inherited_rows += 1

    inheritance = {
        "rule": "same_schedule_same_item_unique_name",
        "mapped_items": sum(
            1
            for catalog in idx.catalog_by_item.values()
            if catalog.get("cn_name") or catalog.get("english_name")
        ),
        "applied_rows": inherited_rows,
        "conflicts": idx.inheritance_conflicts,
    }
    if idx.inheritance_conflicts:
        warnings.append(
            f"排期中有 {len(idx.inheritance_conflicts)} 组货号/品名冲突，相关品名未自动继承，请查看“排期继承检查”"
        )
    export_file = (
        generate_new_rows_excel(new_rows, master_path, export_dir, inheritance)
        if new_rows else ""
    )
    return {
        "ok": True,
        "msg": f"已按新单生成 {len(new_rows)} 行 Excel，原排期未修改" if new_rows else "没有识别到可生成的新单行",
        "modified": 0,
        "new_count": len(new_rows),
        "mod_details": [],
        "new_details": new_details,
        "rows": database_rows,
        "export_file": export_file,
        "warnings": warnings,
        "inheritance": inheritance,
        "mode": "new_order_only",
    }


_UPDATE_COLS = {
    # 来单期不在改单覆盖之列：业务口径是"客户确认下单的邮件日期"（见做排期规则文档），
    # PO 上的 Order Date 只作新单参考值，不覆盖人工填的邮件日。
    COL["customer_po"]: "客PO",
    COL["contract"]: "合同号",
    # 客名不在改单覆盖之列：人工规范写法(如 Target USA)优先于解析值(TARGET)；
    # 新单仍由 _compose_row 填客名。
    COL["version"]: "版本",
    COL["qty"]: "数量",
    COL["inner"]: "内箱",
    COL["outer"]: "外箱",
    COL["inspection_date"]: "验货期",
    COL["ship_date"]: "PO走货期",
    COL["unit_usd"]: "单价/(美金)",
    COL["total_usd"]: "总货价/（美金）",
}


def _update_existing_row(ws, row_no: int, values: dict[int, Any]) -> list[str]:
    changes = []
    qty_or_outer_changed = False
    for col, label in _UPDATE_COLS.items():
        new = values.get(col)
        if new in (None, ""):
            continue
        cell = ws.cell(row_no, col)
        old = cell.value
        if _value_equal(old, new):
            continue
        # 一张 PO 可能带多个客PO(如 "9114023/9114040")；总表按行只记其中一个时不覆盖。
        if col == COL["customer_po"] and "/" in str(new) and _display(old).strip() in str(new).split("/"):
            continue
        cell.value = new
        cell.fill = BLUE_FILL
        changes.append(f"{label} {_display(old)}→{_display(new)}")
        if col in (COL["qty"], COL["outer"]):
            qty_or_outer_changed = True
    if qty_or_outer_changed:
        total_cell = ws.cell(row_no, COL["total_box"])
        qty_cell = ws.cell(row_no, COL["qty"]).value
        outer_cell = ws.cell(row_no, COL["outer"]).value
        if qty_cell not in (None, "") and outer_cell not in (None, "", 0):
            formula = f'=IF(M{row_no}=0,"",K{row_no}/M{row_no})'
            if total_cell.value != formula:
                total_cell.value = formula
                total_cell.fill = BLUE_FILL
                changes.append("总箱数 公式")
    return changes


def generate_new_rows_excel(
    rows: list[dict[int, Any]],
    master_path: str,
    output_dir: str,
    inheritance: dict[str, Any] | None = None,
) -> str:
    src = openpyxl.load_workbook(master_path)
    try:
        src_ws, header_row = _select_master_sheet(src)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "新单数据"
        _copy_header(src_ws, ws, EXPORT_END_COL, header_row)
        _write_rows(src_ws, ws, rows, EXPORT_END_COL, header_row)

        core = wb.create_sheet("仓单PO模版")
        _copy_header(src_ws, core, CORE_END_COL, header_row)
        _write_rows(src_ws, core, rows, CORE_END_COL, header_row)

        _autosize(ws, EXPORT_END_COL)
        _autosize(core, CORE_END_COL)
        if inheritance and (inheritance.get("applied_rows") or inheritance.get("conflicts")):
            check = wb.create_sheet("排期继承检查")
            headers = ("状态", "范围", "货号", "字段", "排期值")
            for col, title in enumerate(headers, 1):
                check.cell(1, col, title).fill = BLUE_FILL
            check.cell(2, 1, "已继承标准品名")
            check.cell(2, 2, "新单")
            check.cell(2, 3, f"{inheritance.get('applied_rows', 0)} 行")
            row_no = 3
            for conflict in inheritance.get("conflicts", []):
                values = (
                    "排期存在冲突，未自动继承",
                    conflict.get("scope", ""),
                    conflict.get("item", ""),
                    conflict.get("field", ""),
                    " / ".join(conflict.get("values", [])),
                )
                for col, value in enumerate(values, 1):
                    check.cell(row_no, col, value)
                row_no += 1
            check.freeze_panes = "A2"
            for col, width in enumerate((26, 18, 20, 18, 60), 1):
                check.column_dimensions[get_column_letter(col)].width = width
        os.makedirs(output_dir, exist_ok=True)
        fname = f"Casdon新单_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.xlsx"
        out = os.path.join(output_dir, fname)
        wb.save(out)
        wb.close()
        return fname
    finally:
        src.close()


def _copy_header(src_ws, dest_ws, end_col: int, header_row: int) -> None:
    for row_no in range(1, header_row + 1):
        if src_ws.row_dimensions[row_no].height:
            dest_ws.row_dimensions[row_no].height = src_ws.row_dimensions[row_no].height
        for col in range(1, end_col + 1):
            src_cell = src_ws.cell(row_no, col)
            dest = dest_ws.cell(row_no, col, src_cell.value)
            if src_cell.has_style:
                dest.font = copy(src_cell.font)
                dest.fill = copy(src_cell.fill)
                dest.border = copy(src_cell.border)
                dest.alignment = copy(src_cell.alignment)
                dest.number_format = src_cell.number_format
                dest.protection = copy(src_cell.protection)
    for col in range(1, end_col + 1):
        width = src_ws.column_dimensions[get_column_letter(col)].width
        if width:
            dest_ws.column_dimensions[get_column_letter(col)].width = width
    for merged in src_ws.merged_cells.ranges:
        if merged.max_row <= header_row and merged.max_col <= end_col:
            dest_ws.merge_cells(str(merged))
    dest_ws.freeze_panes = f"B{header_row + 1}"
    dest_ws.auto_filter.ref = f"A{header_row}:{get_column_letter(end_col)}{header_row}"


def _write_rows(src_ws, ws, rows: list[dict[int, Any]], end_col: int, header_row: int) -> None:
    style_row = header_row
    for candidate in range(header_row + 1, min(src_ws.max_row, header_row + 30) + 1):
        if src_ws.cell(candidate, COL["item"]).value not in (None, ""):
            style_row = candidate
            break
    start_row = header_row + 1
    for offset, data in enumerate(rows, start=start_row):
        if src_ws.row_dimensions[style_row].height:
            ws.row_dimensions[offset].height = src_ws.row_dimensions[style_row].height
        for col in range(1, end_col + 1):
            value = data.get(col, "")
            if col == COL["total_box"] and data.get(COL["qty"]) not in (None, ""):
                value = f'=IF(M{offset}=0,"",K{offset}/M{offset})'
            cell = ws.cell(offset, col, value)
            source = src_ws.cell(style_row, col)
            if source.has_style:
                cell.font = copy(source.font)
                cell.fill = copy(source.fill)
                cell.border = copy(source.border)
                cell.alignment = copy(source.alignment)
                cell.number_format = source.number_format
                cell.protection = copy(source.protection)
            if col in (COL["po_date"], COL["inspection_date"], COL["ship_date"]) and value not in (None, ""):
                cell.number_format = "yyyy/m/d"
            if col in (COL["customer_po"], COL["contract"]):
                cell.number_format = "@"


def _autosize(ws, end_col: int) -> None:
    for col in range(1, end_col + 1):
        letter = get_column_letter(col)
        if ws.column_dimensions[letter].width:
            continue
        ws.column_dimensions[letter].width = 12
