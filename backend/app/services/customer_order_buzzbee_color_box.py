"""Expose proven BuzzBee packaging references in the unified review sheet."""
from collections import defaultdict
from copy import copy

from app.services.customer_order_buzzbee import PACKAGING_PATTERN


HEADER = "彩盒编号"


def _literal(cell):
    return "" if cell.data_type == "f" or cell.value is None else str(cell.value).strip()


def _code(cell):
    text = _literal(cell)
    return text if PACKAGING_PATTERN.fullmatch(text) else ""


def _identity(sheet, number, columns):
    # Keep separators and product variants: SO-A and SOA are different orders.
    values = tuple(_literal(sheet.cell(number, column)).casefold() for column in columns)
    return values if all(values) else None


def _active_end(sheet):
    from app.services.customer_order_unified import _marker_row
    cancelled = _marker_row(sheet)
    shipped = [number for (number, _), cell in sheet._cells.items()
               if number >= 4 and "已走货" in _literal(cell)]
    return min([cancelled, *shipped])


def _write(sheet, number, column, value):
    cell = sheet.cell(number, column)
    if not cell.has_style:
        cell._style = copy(sheet.cell(number, 20)._style)
    cell.number_format = "@"
    cell.value = value


def prepare_review(workbook):
    """Add/reuse the column and fill only proven, empty active-history cells."""
    from app.services import customer_order_huaxing_unified as mapping
    from app.services.customer_order_unified import CustomerOrderUnifiedError, _normalized_header

    review = workbook["正单评审表"]
    matches = [cell.column for cell in review[3] if _normalized_header(cell.value) == HEADER]
    if len(matches) > 1:
        raise CustomerOrderUnifiedError("正单评审表的“彩盒编号”列重复，请先核对表头")
    if matches:
        column = matches[0]
        if column <= 19 or any(review.cell(3, column).coordinate in area
                               for area in review.merged_cells.ranges):
            raise CustomerOrderUnifiedError("正单评审表的“彩盒编号”占用公共字段或合并表头，请先核对表头")
    else:
        # max_column includes body cells and merged ranges even with a blank
        # header; in real templates the blank U header has cancelled-row formulas.
        column = mapping._extension_column(review, HEADER, after=20)
        if review.print_area:
            from openpyxl.worksheet.print_settings import PrintArea
            areas = PrintArea.from_string(review.print_area)
            expanded = []
            for area in areas.ranges:
                if area.min_col <= 20 <= area.max_col:
                    area.max_col = max(area.max_col, column)
                expanded.append(str(area))
            review.print_area = expanded

    sources = defaultdict(list)
    for name in mapping.item_sheets(workbook, "buzzbee"):
        sheet = workbook[name]
        for number in range(4, _active_end(sheet)):
            key = _identity(sheet, number, (5, 8, 7))  # SO, full item, end customer
            if key:
                sources[key].append(_code(sheet.cell(number, 20)))
    remarks = [cell.column for cell in review[3]
               if cell.column > 19 and _normalized_header(cell.value) == "备注"]
    for number in range(4, _active_end(review)):
        target = review.cell(number, column)
        # Formula cells (even formulas displaying "") and manual values are owned
        # by the user. Whitespace is also an existing value.
        if target.value not in (None, ""):
            continue
        if not (_literal(review.cell(number, 7)) and
                (_literal(review.cell(number, 4)) or _literal(review.cell(number, 3)))):
            continue
        key = _identity(review, number, (4, 7, 6))
        candidates = sources.get(key, [])
        note = _code(review.cell(number, remarks[0])) if len(remarks) == 1 else ""
        if len(candidates) > 1:
            continue
        item_code = candidates[0] if candidates else ""
        if item_code and note and item_code != note:
            continue
        value = item_code or note
        if value:
            _write(review, number, column, value)
    return column


def write_new_row(workbook, column, review_row, item_sheet, item_row):
    """A newly generated order links to its exact ITEM row, never another PO."""
    value = None
    if _code(item_sheet.cell(item_row, 20)):
        name = item_sheet.title.replace("'", "''")
        ref = f"'{name}'!T{item_row}"
        value = f'=IF(LEN({ref})=0,"",{ref})'
    _write(workbook["正单评审表"], review_row, column, value)
