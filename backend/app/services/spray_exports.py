from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from app.models import spray_production as m
from app.services.spray_production import serial

EXPORTS = {
    'orders': (m.SprayOrder, [('document_no','工单号'),('customer','委托客户'),('due_date','交付日期'),('status','状态')], False),
    'shipments': (m.SprayShipment, [('document_no','单据号'),('customer','客户'),('business_date','业务日期'),('kind','送货/退货'),('reason','说明')], False),
    'returnables': (m.SprayReturnable, [('customer','客户'),('name','容器'),('quantity','发出正数/收回负数'),('unit','单位'),('document_no','关联单据')], False),
    'settlements': (m.SpraySettlement, [('customer','客户'),('period','月份'),('currency','币种'),('amount','净结算金额'),('status','状态')], True),
    'shipment-lines': (m.SprayShipmentLine, [('shipment_id','送货记录引用'),('batch_id','批次引用'),('quantity','净件数'),('price','每件单价'),('currency','币种'),('amount','金额'),('price_reference','采用价格依据')], True),
    'wages': (m.SprayReportLine, [('report_id','班次记录引用'),('task_id','任务引用'),('regular_qty','正班数量'),('overtime_qty','加班数量'),('regular_hours','正班小时'),('overtime_hours','加班小时'),('wage_status','工资状态')], True),
}


def workbook_bytes(db, factory, kind):
    cls, columns, _ = EXPORTS[kind]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '喷油业务明细'
    names = {'huaxing':'华兴','huakang-a':'华康 A','huakang-b':'华康 B','huadeng':'华登'}
    sheet.append([names[factory] + ' · 喷油部生产管理'])
    sheet.append([label for _, label in columns])
    for cell in sheet[2]:
        cell.fill = PatternFill('solid', fgColor='006559'); cell.font = Font(color='FFFFFF', bold=True)
    query = select(cls).where(cls.factory_id == factory).order_by(cls.created_at, cls.id)
    for item in db.scalars(query.execution_options(yield_per=500)):
        values = serial(item)
        sheet.append([values.get(key, '') for key, _ in columns])
        for cell in sheet[sheet.max_row]:
            if cell.data_type == 'f': cell.data_type = 's'
    sheet.freeze_panes = 'B3'
    sheet.auto_filter.ref = f'A2:{sheet.cell(2,len(columns)).column_letter}{sheet.max_row}'
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24
    buffer = BytesIO(); workbook.save(buffer); return buffer.getvalue()
