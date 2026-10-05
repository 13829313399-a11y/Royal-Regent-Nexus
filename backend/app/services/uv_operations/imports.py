"""Explicit, versioned templates. Preview never changes a formal ledger."""
import csv
from hashlib import sha256
from io import BytesIO, StringIO
import json
from pathlib import PurePath
from zipfile import ZipFile, BadZipFile
from uuid import uuid4
from itertools import islice
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.exc import IntegrityError
from datetime import UTC,datetime,timedelta
from app.models import uv_operations as m
from app.schemas import uv_operations as s
from . import common as c, authz as a, production as p, reports

MAX_BYTES = 16*1024*1024
TEMPLATES = {
    'reference-efficiency-v1':(s.ReferenceEfficiency,('cost_read','cost_write'),lambda db,user,body:p.create_master(db,user,body,m.UvOpsReferenceEfficiency)),
    "products-v1": (s.ProductCreate, ("master_write",), lambda db, user, body: p.create_master(db, user, body, m.UvOpsProduct)),
    "processes-v1": (s.ProcessCreate, ("master_write",), lambda db, user, body: p.create_master(db, user, body, m.UvOpsProcessVersion)),
    "demands-v1": (s.DemandCreate, ("plan_write",), p.create_demand),
    "manual-production-v1": (s.ProductionConfirm, ("production_write",), p.confirm),
    "expenses-v1": (s.ExpenseCreate, ("cost_read", "cost_write"), reports.expense),
}


def read_file(name, content):
    c.require(name and '/' not in name and '\\' not in name and '..' not in name, "unsafe_filename", "文件名不能包含路径", 422)
    c.require(0 < len(content) <= MAX_BYTES, "import_size", "导入文件为空或超过 16 MB", 422)
    suffix = PurePath(name).suffix.lower()
    warnings = []
    if suffix == ".csv":
        try:
            values = list(islice(csv.reader(StringIO(content.decode("utf-8-sig"))),10002))
            c.require(all(len(row)<=128 for row in values),'import_columns','导入最多 128 列',422)
        except (UnicodeError, csv.Error):
            raise c.DomainError("csv_encoding", "CSV 需使用 UTF-8 编码", 422) from None
        sheets = [("CSV", values)]
    elif suffix == ".xlsx":
        try:
            with ZipFile(BytesIO(content)) as archive:
                entries = archive.infolist()
                c.require(len(entries) <= 500 and sum(x.file_size for x in entries) <= 64*1024*1024 and all(x.file_size <= max(x.compress_size,1)*200 for x in entries), "unsafe_archive", "工作簿解压大小或压缩比超限", 422)
                c.require(not any('..' in x.filename or x.filename.startswith('/') for x in entries), "unsafe_archive", "工作簿包含不安全路径", 422)
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
            sheets=[]
            for sheet in workbook:
                c.require((sheet.max_row or 0)<=10001 and (sheet.max_column or 0)<=128,'worksheet_dimensions','工作表声明范围超过 10000 行或 128 列，请清理多余格式',422)
                rows=list(islice(sheet.iter_rows(values_only=True,max_col=sheet.max_column or 128),10002))
                sheets.append((sheet.title,rows))
            workbook.close()
            warnings.append("公式仅读取缓存值；缺缓存会报字段错误，宏与外部链接不执行")
        except (BadZipFile, OSError, ValueError, KeyError):
            raise c.DomainError("invalid_workbook", "无法识别 XLSX 工作簿", 422) from None
    else:
        raise c.DomainError("import_format", "仅接收 UTF-8 CSV 或 XLSX；不接受宏工作簿", 422)
    c.require(sum(len(rows) for _, rows in sheets) <= 10001, "import_row_limit", "单批最多 10000 行，请按明确批次拆分导入", 422)
    result = []
    for sheet, rows in sheets:
        if not rows:
            continue
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        c.require(all(headers) and len(headers) == len(set(headers)), "import_headers", "表头不能为空或重复", 422)
        for index, values in enumerate(rows[1:], start=2):
            if all(value is None or value == "" for value in values):
                continue
            result.append(dict(sheet=sheet, row=index, headers=headers, values={key:value for key,value in zip(headers, values) if value is not None and value != ""}))
    return result, warnings


def preview(db, user, template, name, content, field_mapping=None, units_confirmed=False):
    c.require(template in TEMPLATES, "import_template", "请选择明确的导入模板版本", 422)
    schema, permissions, _ = TEMPLATES[template]
    a.authorize(user, m.FACTORY, "import", *permissions)
    c.require(units_confirmed, 'units_confirmation', '请先确认模板单位：pcs、mm、秒及明确币种；不自动猜测单位', 422)
    field_mapping = field_mapping or {}
    c.require(isinstance(field_mapping,dict) and len(field_mapping)<=128 and all(isinstance(k,str) and isinstance(v,str) and k in schema.model_fields and k not in {'factory_id','operation_id','reason'} and len(v)<=128 for k,v in field_mapping.items()), 'field_mapping', '映射应为模板字段到原始表头的对象，且只包含有效业务字段', 422)
    c.require(len(set(field_mapping.values()))==len(field_mapping), 'field_mapping', '一个原始列不能重复映射多个字段', 422)
    c.require(0<len(content)<=MAX_BYTES,'import_size','导入文件为空或超过 16 MB',422)
    job=c.add(db,m.UvOpsImportJob,user,actor_id=user.id,template=template,file_hash=sha256(content).hexdigest(),field_mapping=field_mapping,units_confirmed=True,rows=[],errors=[],source_name=name,source_content=content,status='queued')
    if len(content)<=128*1024:
        build_preview(db,user,job)
    db.commit()
    return summary(job)


def summary(job):
    return dict(id=job.id,version=job.version,template=job.template,field_mapping=job.field_mapping,units_confirmed=job.units_confirmed,row_count=len(job.rows)+len(job.errors),valid_count=len(job.rows),errors=job.errors,warnings=job.warnings,preview=job.rows[:20],preview_truncated=len(job.rows)>20,status=job.status,results=job.results)


def build_preview(db,user,job):
    template,name,content=job.template,job.source_name,job.source_content
    schema,permissions,handler=TEMPLATES[template]
    a.authorize(user,m.FACTORY,'import',*permissions)
    rows, warnings = read_file(name, content)
    file_hash = sha256(content).hexdigest()
    errors, normalized = [], []
    for row in rows:
        values = dict(row["values"])
        try:
            if job.field_mapping:
                missing = set(job.field_mapping.values()) - set(row['headers'])
                c.require(not missing, 'mapping_source_missing', '映射引用的原始列不存在', 422)
                mapped = {target:values[source] for target,source in job.field_mapping.items() if source in values}
                values = {key:value for key,value in values.items() if key not in job.field_mapping.values()} | mapped
            for key, field in schema.model_json_schema()["properties"].items():
                if key in values and isinstance(values[key],float):
                    # Excel numeric cells are normalized once into decimal text;
                    # the canonical business API never accepts binary floats.
                    values[key]=str(values[key])
                if key in values and field.get("type") == "integer" and isinstance(values[key], str):
                    c.require(values[key].lstrip('-').isdigit(), "integer_required", "件数及版本必须是整数", 422)
                    values[key] = int(values[key])
            c.require(values.get("factory_id", m.FACTORY) == m.FACTORY, "cross_factory_row", "文件包含非华康 A 数据", 422)
            command = schema.model_validate(dict(factory_id=m.FACTORY, operation_id=uuid4().hex, expected_version=0) | values)
            identity = c.digest(dict(file_hash=file_hash, template=template, sheet=row["sheet"], row=row["row"], identity=values.get("code", values.get("batch_id")), units="template-v1",mapping=job.field_mapping))
            normalized.append(dict(**row, identity=identity, command=command.model_dump(mode="json")))
        except (ValidationError, c.DomainError) as error:
            errors.append(dict(sheet=row["sheet"], row=row["row"], message=error.body["message"] if isinstance(error, c.DomainError) else "字段类型、单位或必填项不符合模板", fields=[str(x["loc"][0]) for x in error.errors()] if isinstance(error, ValidationError) else []))
    # Execute the exact domain validators against a rollback-only savepoint.
    # Earlier valid rows are visible to later rows, so preview catches duplicates,
    # batch versions and cumulative balances without changing the formal ledger.
    checked=[]
    outer=db.begin_nested()
    try:
        for row in normalized:
            nested=db.begin_nested()
            try:
                existing=db.scalar(c.query(m.UvOpsImportRow).where(m.UvOpsImportRow.identity==row['identity']))
                if not existing:
                    handler(db,user,schema.model_validate(row['command']))
                    db.flush()
                checked.append(row|dict(disposition='duplicate' if existing else 'new'))
                nested.commit()
            except (c.DomainError,IntegrityError,ValidationError) as error:
                nested.rollback()
                code=error.body['code'] if isinstance(error,c.DomainError) else 'row_constraint'
                errors.append(dict(sheet=row['sheet'],row=row['row'],code=code,message=error.body['message'] if isinstance(error,c.DomainError) else '关联、版本或唯一性校验失败',fields=[],disposition='adjustment_draft' if code=='period_closed' else 'error'))
    finally:
        outer.rollback()
    job.rows,job.errors,job.warnings=checked,errors,warnings
    job.status='adjustment_required' if any(x.get('disposition')=='adjustment_draft' for x in errors) else 'preview'
    job.lease_until=None
    c.touch(job)
    return job


def commit(db, user, entity_id, body):
    job = c.get(db, m.UvOpsImportJob, entity_id, lock=True, version=body.expected_version)
    c.require(job.actor_id == user.id, "import_owner", "只能确认本人上传的导入预览", 403)
    schema, permissions, handler = TEMPLATES[job.template]
    a.authorize(user, m.FACTORY, "import", *permissions)
    c.require(not job.errors and job.status == "preview", "import_invalid", "请先修正错误并生成新的导入预览")
    imported, duplicates, results = [], 0, []
    for row in job.rows:
        existing = db.scalar(c.query(m.UvOpsImportRow).where(m.UvOpsImportRow.identity == row["identity"]))
        if existing:
            duplicates += 1
            results.append(dict(sheet=row['sheet'],row=row['row'],status='duplicate',entity_id=existing.entity_id))
            continue
        try:
            result = handler(db, user, schema.model_validate(row["command"]))
            db.flush()
        except (c.DomainError,IntegrityError) as error:
            raise c.DomainError('import_row_conflict',f"工作表 {row['sheet']} 第 {row['row']} 行已变化，整批未入账，请重新校验",409) from None
        entity_id = result.get("id") or result.get("entry", {}).get("id")
        c.add(db, m.UvOpsImportRow, user, identity=row["identity"], job_id=job.id, entity_id=entity_id)
        imported.append(entity_id)
        results.append(dict(sheet=row['sheet'],row=row['row'],status='imported',entity_id=entity_id))
    job.status = "committed"
    job.results=results
    c.touch(job)
    return dict(id=job.id, imported=len(imported), duplicates=duplicates, entity_ids=imported, status=job.status,results=results)


def revalidate(db,user,entity_id,body):
    job=c.get(db,m.UvOpsImportJob,entity_id,lock=True,version=body.expected_version)
    c.require(job.actor_id==user.id,'import_owner','只能重新校验本人导入',403)
    c.require(job.status in {'preview','adjustment_required','failed'},'import_state','当前状态不可重新校验')
    return summary(build_preview(db,user,job))


def process_one():
    from app.db import SessionLocal
    from app.models.auth import AuthUser
    from app.services.auth import build_auth_context
    with SessionLocal() as db:
        now=c.ts(datetime.now(UTC))
        job=db.scalar(c.query(m.UvOpsImportJob).where((m.UvOpsImportJob.status=='queued')|((m.UvOpsImportJob.status=='processing')&(m.UvOpsImportJob.lease_until<now))).order_by(m.UvOpsImportJob.created_at).with_for_update(skip_locked=True))
        if job is None:return
        job.status,job.lease_until='processing',c.ts(datetime.now(UTC)+timedelta(minutes=10))
        identifier=job.id
        db.commit()
    with SessionLocal() as db:
        job=c.get(db,m.UvOpsImportJob,identifier,lock=True)
        try:
            user=db.get(AuthUser,job.actor_id)
            c.require(user is not None and user.status=='active','permission_denied','导入账号已失效',403)
            build_preview(db,build_auth_context(db,user),job)
            db.commit()
        except Exception as error:
            db.rollback()
            job=c.get(db,m.UvOpsImportJob,identifier,lock=True)
            job.status,job.lease_until='failed',None
            job.errors=[dict(code=error.body['code'] if isinstance(error,c.DomainError) else 'preview_failed',message=error.body['message'] if isinstance(error,c.DomainError) else '预览失败，请检查文件后重新校验')]
            db.commit()
