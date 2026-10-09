"""Durable export jobs, complete database queries, authorization at download."""
import asyncio
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
import json
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter
from itertools import chain
from openpyxl.styles import Font, PatternFill
from sqlalchemy import Numeric, select, func
from app.db import SessionLocal
from app.models import uv_operations as m
from app.models.auth import AuthUser
from app.services.auth import build_auth_context
from . import common as c, authz as a, reports

MODELS = {"tasks":m.UvOpsTask, "schedule":m.UvOpsScheduleBlock, "production":m.UvOpsProductionEntry, "handovers":m.UvOpsHandover, "ink":m.UvOpsInkMovement, "wages":m.UvOpsWageAccrual}


def permissions(kind):
    return ("export", "payroll_read") if kind == "wages" else ("export", "cost_read", "payroll_read") if kind in {"daily", "monthly"} else ("export",)


def create(db, user, body):
    a.authorize(user, body.factory_id, *permissions(body.kind))
    c.require(body.kind not in {'daily','monthly'} or not body.selected_ids, 'report_selection', '日报月报仅支持完整日期范围，不能按记录 ID 导出局部汇总', 422)
    start, end = date.fromisoformat(body.start_date), date.fromisoformat(body.end_date)
    c.require(start <= end and (end-start).days <= 366, "export_window", "导出区间最多 366 天", 422)
    row = c.add(db, m.UvOpsExportJob, user, actor_id=user.id, kind=body.kind, filters=c.values(body, "kind"), permissions=list(permissions(body.kind)))
    return c.record(row)


def selection_statement(job):
    filters = job.filters
    model = MODELS[job.kind]
    statement = c.query(model, history=True).order_by(model.created_at, model.id)
    if filters["selected_ids"]:
        statement = statement.where(model.id.in_(filters["selected_ids"]))
    else:
        field = getattr(model, "business_date", model.created_at)
        if job.kind == 'schedule':
            statement = statement.where(model.start_at<c.ts(str(date.fromisoformat(filters['end_date'])+timedelta(days=1))+'T00:00:00+08:00'),model.end_at>c.ts(filters['start_date']+'T00:00:00+08:00'))
        elif hasattr(model, "business_date"):
            statement = statement.where(field.between(filters["start_date"], filters["end_date"]))
        else:
            statement = statement.where(field >= c.ts(filters["start_date"]+"T00:00:00+08:00"), field < c.ts(str(date.fromisoformat(filters["end_date"])+timedelta(days=1))+"T00:00:00+08:00"))
    return statement


def rows(db, user, job, report=None):
    filters = job.filters
    if job.kind in {"daily", "monthly"}:
        c.require(not filters['selected_ids'], 'report_selection', '日报月报不支持选中记录范围', 422)
        value = report if report is not None else reports.report(db, filters["start_date"], filters["end_date"])
        return value["details"], {'cost_revenue'}, value["coverage"]
    model = MODELS[job.kind]
    statement = selection_statement(job)
    result = (a.project(c.record(row), user) for row in db.scalars(statement.execution_options(yield_per=500)))
    if filters["selected_ids"]:
        c.require(db.scalar(select(func.count()).select_from(statement.subquery())) == len(set(filters["selected_ids"])), "export_selection_missing", "部分选中记录不存在或不属于当前厂区", 422)
    numeric = {column.name for column in model.__table__.columns if isinstance(column.type, Numeric)}
    return result, numeric, dict(state="records", missing_cost_records=None)


def append_sheet(workbook, title, records, numeric):
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = 'A2'
    sheet.print_title_rows = '1:1'
    iterator = iter(records)
    first = next(iterator, None)
    headers = list(first) if first else ['无匹配记录']
    def append(values):
        cells = []
        for value in values:
            if isinstance(value, (dict,list)): value = json.dumps(value,ensure_ascii=False)
            cell = WriteOnlyCell(sheet,value)
            if isinstance(value,str): cell.data_type = 's'
            cells.append(cell)
        sheet.append(cells)
    append(headers)
    for row in chain([first],iterator) if first else ():
        append([Decimal(row[key]) if key in numeric and row.get(key) is not None else row.get(key) for key in headers])
    for index in range(1,len(headers)+1): sheet.column_dimensions[get_column_letter(index)].width=22


def build(db, user, job):
    a.authorize(user, m.FACTORY, *job.permissions)
    report = reports.report(db,job.filters['start_date'],job.filters['end_date']) if job.kind in {'daily','monthly'} else None
    data, numeric, coverage = rows(db, user, job, report)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("UV记录")
    iterator=iter(data)
    first=next(iterator,None)
    headers=list(first) if first else ["无匹配记录"]
    sheet.freeze_panes="A2"
    sheet.print_title_rows="1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.orientation="landscape"
    sheet.page_setup.fitToWidth=1
    sheet.oddFooter.center.text="核数人：________  复核人：________  签收人：________"
    sheet.oddFooter.right.text='第 &P 页 / 共 &N 页'
    for index in range(1,len(headers)+1):
        sheet.column_dimensions[get_column_letter(index)].width=20
    title=[]
    for value in headers:
        cell=WriteOnlyCell(sheet,value)
        cell.data_type='s'
        cell.font=Font(bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='095C55')
        title.append(cell)
    sheet.append(title)
    count=0
    for row in chain([first],iterator) if first else ():
        values=[]
        for key in headers:
            value=row.get(key)
            if value is not None and key in numeric:
                value=Decimal(value)
            elif isinstance(value,(dict,list)):
                value=json.dumps(value,ensure_ascii=False)
            cell=WriteOnlyCell(sheet,value)
            if isinstance(value,str): cell.data_type='s'
            values.append(cell)
        sheet.append(values)
        count+=1
    sheet.auto_filter.ref=f'A1:{get_column_letter(len(headers))}{count+1}'
    if job.kind in {'production','handovers','wages'}:
        selected = selection_statement(job).with_only_columns(MODELS[job.kind].id).order_by(None)
        if job.kind == 'production':
            model, title, statement = m.UvOpsQualityEntry, '品质处置明细', c.query(m.UvOpsQualityEntry, history=True)
            statement = statement.where(model.production_entry_id.in_(selected)) if job.filters['selected_ids'] else statement.where(model.business_date.between(job.filters['start_date'],job.filters['end_date']))
        elif job.kind == 'handovers':
            model, title, statement = m.UvOpsHandoverEvent, '签收拒收退回', c.query(m.UvOpsHandoverEvent)
            statement = statement.where(model.handover_id.in_(selected)) if job.filters['selected_ids'] else statement.where(model.business_date.between(job.filters['start_date'],job.filters['end_date']))
        else:
            model, title, statement = m.UvOpsWageAllocation, '员工工资分配', c.query(m.UvOpsWageAllocation, history=True).where(m.UvOpsWageAllocation.accrual_id.in_(selected))
        linked = (a.project(c.record(row), user) for row in db.scalars(statement.order_by(model.created_at,model.id).execution_options(yield_per=500)))
        append_sheet(workbook,title,linked,{column.name for column in model.__table__.columns if isinstance(column.type,Numeric)})
        if job.kind in {'production','handovers'}:
            source_model,source_field,source_title = (m.UvOpsProductionEntry,model.production_entry_id,'品质来源核数') if job.kind=='production' else (m.UvOpsHandover,model.handover_id,'关联交接原单')
            source_rows = db.scalars(c.query(source_model).where(source_model.id.in_(statement.with_only_columns(source_field).order_by(None))))
            append_sheet(workbook,source_title,(a.project(c.record(row),user) for row in source_rows),{column.name for column in source_model.__table__.columns if isinstance(column.type,Numeric)})
        if job.kind in {'production','wages'}:
            kind = 'quality' if job.kind == 'production' else 'wage'
            targets = statement.with_only_columns(model.id).order_by(None) if kind == 'quality' else selected
            receipts = db.scalars(c.query(m.UvOpsReversal).where(m.UvOpsReversal.kind==kind,m.UvOpsReversal.entity_id.in_(targets)))
            append_sheet(workbook,'撤销依据',(c.record(row) for row in receipts),set())
    if job.kind in {'daily','monthly'}:
        totals=workbook.create_sheet('按币种核算')
        columns=['currency','cost_revenue','cost_amount','payroll_total','cost_total','cost_margin']
        totals.append(columns)
        for currency,values in report['cost_by_currency'].items():
            totals.append([currency]+[Decimal(values[key]) if values.get(key) is not None else None for key in columns[1:]])
        totals.freeze_panes='A2'
        totals.print_title_rows='1:1'
        missing=workbook.create_sheet('数据缺项')
        missing.append(['task_id','shift_id','code'])
        for item in report['cost_missing']:
            missing.append([item.get(key) for key in ('task_id','shift_id','code')])
        warnings=workbook.create_sheet('警告与计算口径')
        for key,value in [('coverage',report['coverage']),('warnings',report['warnings']),('report_basis',report.get('report_basis','current_open_ledger')),('totals',report['totals']),('yield_ratio',report['yield_ratio'])]:
            warnings.append([key,json.dumps(value,ensure_ascii=False)])
    meta = workbook.create_sheet("口径与范围")
    for key, value in dict(factory_id=m.FACTORY, timezone="Asia/Shanghai", as_of=m.now(), template_version="uv-export-v2", linked_scope="selected: complete linked history; date range: quality and handover events by their own business_date, source sheets retain original dates", scope="selected" if job.filters["selected_ids"] else "filtered_database", row_count=count, filters=job.filters, coverage=coverage, quantity_unit="pcs (unless column explicitly says board/ml/seconds)", money="currency per row; decimal numeric cells", data_mode=db.scalar(c.query(m.UvOpsSettings)).data_mode).items():
        meta.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value,(dict,list)) else value])
    output = BytesIO()
    workbook.save(output)
    # Artifact is immutable once ready. If sensitive read access is subsequently
    # removed, downloading this artifact is denied; create a new redacted job.
    job.filters = dict(job.filters, _built_read_permissions=[permission for permission in ('cost_read','payroll_read') if a.allowed(user,permission)])
    job.artifact, job.row_count, job.status, job.lease_until = output.getvalue(), count, "ready", None
    c.touch(job)
    return job


def process_one():
    with SessionLocal() as db:
        if not __import__('sqlalchemy').inspect(db.bind).has_table("uv_ops_export_jobs"):
            return
        now = c.ts(datetime.now(UTC))
        job = db.scalar(c.query(m.UvOpsExportJob).where((m.UvOpsExportJob.status == "queued") | ((m.UvOpsExportJob.status == "running") & (m.UvOpsExportJob.lease_until < now))).order_by(m.UvOpsExportJob.created_at).with_for_update(skip_locked=True))
        if job is None:
            return
        job.status, job.lease_until = "running", c.ts(datetime.now(UTC)+timedelta(minutes=10))
        job_id = job.id
        db.commit()
    with SessionLocal() as db:
        if db.bind.dialect.name == 'postgresql':
            db.connection(execution_options={'isolation_level':'REPEATABLE READ'})
        job = c.get(db, m.UvOpsExportJob, job_id, lock=True)
        try:
            account = db.get(AuthUser, job.actor_id)
            c.require(account is not None and account.status == "active", "permission_denied", "导出账号已失效", 403)
            build(db, build_auth_context(db, account), job)
            db.commit()
        except Exception as error:
            db.rollback()
            job = c.get(db, m.UvOpsExportJob, job_id, lock=True)
            job.status, job.error, job.lease_until = "failed", error.body["code"] if isinstance(error, c.DomainError) else "export_failed", None
            db.commit()


async def worker():
    while True:
        try:
            from . import imports
            await asyncio.to_thread(imports.process_one)
            await asyncio.to_thread(process_one)
        except Exception:
            import logging
            logging.getLogger("uv_ops").warning("export_worker_unavailable")
        await asyncio.sleep(2)
