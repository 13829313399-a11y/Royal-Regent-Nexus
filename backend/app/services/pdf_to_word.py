from __future__ import annotations

import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import pdfplumber
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt
from PIL import ImageStat
from pypdf import PdfReader

from app.services import pdf_to_excel as pdf_extraction

MAX_PDF_PAGES = 80
_TABLE_SETTINGS = {
    "vertical_strategy": "lines",
    "horizontal_strategy": "lines",
    "snap_tolerance": 4,
    "join_tolerance": 4,
    "intersection_tolerance": 5,
}


class PdfToWordConversionError(ValueError):
    pass


@dataclass(frozen=True)
class PdfToWordResult:
    content: bytes
    page_count: int
    table_count: int
    image_count: int
    text_page_count: int
    ocr_page_count: int
    output_file_name: str


@dataclass(frozen=True)
class _PageImage:
    content: bytes
    name: str
    page_number: int
    display_width: float
    page_width: float


@dataclass(frozen=True)
class _PageBlock:
    top: float
    kind: str
    value: str | list[list[str]] | _PageImage


def _clean_text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"[ \t]+", " ", str(value).replace("\x00", " ")).strip()


def _normalize_table(rows: list[list[object]]) -> list[list[str]]:
    normalized = [[_clean_text(cell) for cell in row] for row in rows]
    normalized = [row for row in normalized if any(row)]
    if not normalized:
        return []
    width = max(len(row) for row in normalized)
    return [row + [""] * (width - len(row)) for row in normalized]


def _join_words(words: list[dict]) -> str:
    text = ""
    for word in words:
        token = _clean_text(word.get("text", ""))
        if not token:
            continue
        if not text:
            text = token
        elif token[0] in ",.;:!?%)]}>，。：；、）》】" or text[-1] in "([{</《【":
            text += token
        elif pdf_extraction._is_cjk_character(text[-1]) and pdf_extraction._is_cjk_character(token[0]):
            text += token
        else:
            text += f" {token}"
    return text


def _word_inside_table(word: dict, table_boxes: list[tuple[float, float, float, float]]) -> bool:
    try:
        center_x = (float(word["x0"]) + float(word["x1"])) / 2
        center_y = (float(word["top"]) + float(word["bottom"])) / 2
    except (KeyError, TypeError, ValueError):
        return False
    return any(x0 <= center_x <= x1 and top <= center_y <= bottom for x0, top, x1, bottom in table_boxes)


def _paragraph_blocks(words: list[dict], table_boxes: list[tuple[float, float, float, float]]) -> list[_PageBlock]:
    usable = [
        word
        for word in words
        if _clean_text(word.get("text", "")) and not _word_inside_table(word, table_boxes)
    ]
    lines: list[list[dict]] = []
    line_tops: list[float] = []
    for word in sorted(usable, key=lambda item: (float(item.get("top", 0)), float(item.get("x0", 0)))):
        top = float(word.get("top", 0))
        if lines and abs(top - line_tops[-1]) <= 3.5:
            lines[-1].append(word)
            line_tops[-1] = sum(float(item.get("top", 0)) for item in lines[-1]) / len(lines[-1])
        else:
            lines.append([word])
            line_tops.append(top)

    blocks: list[_PageBlock] = []
    for top, line in zip(line_tops, lines):
        text = _join_words(sorted(line, key=lambda item: float(item.get("x0", 0))))
        if text:
            blocks.append(_PageBlock(top=top, kind="paragraph", value=text))
    return blocks


def _extract_native_blocks(
    page,
) -> tuple[
    list[_PageBlock],
    int,
    str,
    list[tuple[float, float, float, float]],
]:
    native_text = page.extract_text(x_tolerance=3, y_tolerance=3) or ""
    tables: list[tuple[tuple[float, float, float, float], list[list[str]]]] = []
    try:
        for table in page.find_tables(table_settings=_TABLE_SETTINGS) or []:
            rows = _normalize_table(table.extract() or [])
            if rows:
                tables.append((tuple(float(value) for value in table.bbox), rows))
    except Exception:
        tables = []

    table_boxes = [bbox for bbox, _rows in tables]
    try:
        words = page.extract_words(x_tolerance=2, y_tolerance=2, keep_blank_chars=False) or []
    except Exception:
        words = []

    blocks = _paragraph_blocks(words, table_boxes)
    blocks.extend(
        _PageBlock(top=bbox[1], kind="table", value=rows)
        for bbox, rows in tables
    )
    blocks.sort(key=lambda block: block.top)

    if not blocks and native_text.strip():
        blocks = [
            _PageBlock(top=float(index), kind="paragraph", value=line.strip())
            for index, line in enumerate(native_text.splitlines())
            if line.strip()
        ]
    return blocks, len(tables), native_text, table_boxes


def _image_is_meaningful(image, *, display_width: float, display_height: float) -> bool:
    if display_width < 8 or display_height < 8 or display_width * display_height < 250:
        return False

    try:
        pil_image = image.image
        if pil_image.width < 16 or pil_image.height < 16:
            return False
        preview = pil_image.convert("RGB")
        preview.thumbnail((128, 128))
        return max(ImageStat.Stat(preview).stddev) >= 1
    except Exception:
        return False


def _image_as_png(image) -> bytes:
    pil_image = image.image
    if pil_image.mode not in {"RGB", "RGBA", "L", "LA"}:
        pil_image = pil_image.convert("RGBA" if "transparency" in pil_image.info else "RGB")
    output = BytesIO()
    pil_image.save(output, format="PNG")
    return output.getvalue()


def _extract_image_blocks(page, reader_page, *, page_number: int) -> list[_PageBlock]:
    if reader_page is None:
        return []

    try:
        images_by_name = {Path(image.name).stem: image for image in reader_page.images}
    except Exception:
        return []

    # A single PDF image object can be painted hundreds of times as a tiny rule or
    # texture. Keep only its largest appearance on the page so those drawing aids
    # do not become hundreds of separate Word pictures.
    largest_appearances: dict[str, tuple[float, dict]] = {}
    for raw_image in getattr(page, "images", []) or []:
        name = str(raw_image.get("name", ""))
        try:
            display_width = float(raw_image.get("width", 0))
            display_height = float(raw_image.get("height", 0))
        except (TypeError, ValueError):
            continue
        area = display_width * display_height
        previous = largest_appearances.get(name)
        if previous is None or area > previous[0]:
            largest_appearances[name] = (area, raw_image)

    blocks: list[_PageBlock] = []
    for name, (_area, raw_image) in largest_appearances.items():
        image = images_by_name.get(name)
        if image is None:
            continue
        try:
            display_width = float(raw_image.get("width", 0))
            display_height = float(raw_image.get("height", 0))
            top = float(raw_image.get("top", 0))
        except (TypeError, ValueError):
            continue
        if not _image_is_meaningful(
            image,
            display_width=display_width,
            display_height=display_height,
        ):
            continue
        try:
            content = _image_as_png(image)
        except Exception:
            continue
        blocks.append(
            _PageBlock(
                top=top,
                kind="image",
                value=_PageImage(
                    content=content,
                    name=name,
                    page_number=page_number,
                    display_width=display_width,
                    page_width=float(getattr(page, "width", 0) or 0),
                ),
            ),
        )
    return blocks


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _configure_document(document: Document) -> None:
    for section in document.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.6)
        section.right_margin = Cm(1.6)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def _append_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def _append_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(rows):
        for column_index in range(column_count):
            value = row[column_index] if column_index < len(row) else ""
            cell = table.cell(row_index, column_index)
            cell.text = value
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                    run.bold = row_index == 0
            if row_index == 0:
                _set_cell_shading(cell, "DFF5F2")
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def _append_image(document: Document, image: _PageImage) -> None:
    section = document.sections[-1]
    available_width = int(section.page_width - section.left_margin - section.right_margin)
    if image.page_width > 0:
        width_ratio = min(max(image.display_width / image.page_width, 0.12), 1.0)
    else:
        width_ratio = 1.0
    target_width = Emu(max(int(available_width * width_ratio), int(Cm(3))))

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run()
    shape = run.add_picture(BytesIO(image.content), width=target_width)
    shape._inline.docPr.set("title", image.name)
    shape._inline.docPr.set("descr", f"PDF 第 {image.page_number} 页图片")


def _output_file_name(source_file_name: str) -> str:
    source_name = Path(source_file_name or "PDF文件.pdf").name
    stem = re.sub(r'[\\/:*?"<>|]+', "_", Path(source_name).stem).strip(" .") or "PDF文件"
    return f"{stem}_转换结果.docx"


def convert_pdf_to_word(
    pdf_bytes: bytes,
    source_file_name: str,
    *,
    page_text_overrides: Mapping[
        int,
        Sequence[tuple[float, float, str]],
    ]
    | None = None,
) -> PdfToWordResult:
    try:
        pdf = pdfplumber.open(BytesIO(pdf_bytes))
    except Exception as exc:
        raise PdfToWordConversionError("PDF 无法读取；请确认文件未加密、未损坏。") from exc

    document = Document()
    _configure_document(document)
    document.core_properties.title = Path(source_file_name or "PDF文件.pdf").name
    table_count = 0
    image_count = 0
    text_page_count = 0
    ocr_page_count = 0
    converted_page_count = 0

    try:
        image_pdf = PdfReader(BytesIO(pdf_bytes))
    except Exception:
        image_pdf = None

    try:
        page_count = len(pdf.pages)
        if page_count == 0:
            raise PdfToWordConversionError("PDF 中没有可转换的页面。")
        if page_count > MAX_PDF_PAGES:
            raise PdfToWordConversionError(f"单次最多转换 {MAX_PDF_PAGES} 页 PDF。")

        for page_index, page in enumerate(pdf.pages):
            blocks, page_table_count, native_text, table_boxes = _extract_native_blocks(
                page
            )
            if pdf_extraction._native_text_is_unreliable(native_text):
                blocks = []
                page_table_count = 0
                table_boxes = []
            used_native_content = bool(blocks)

            used_ocr = False
            page_override = (
                page_text_overrides.get(page_index + 1)
                if page_text_overrides is not None
                else None
            )
            if page_override:
                overrides = [
                    (float(top), float(bottom), text)
                    for top, bottom, text in page_override
                    if text.strip()
                ]
                full_page_text = any(
                    top <= 1 and bottom >= float(page.height) - 1
                    for top, bottom, _text in overrides
                )
                if full_page_text:
                    blocks = [
                        _PageBlock(top=top, kind="paragraph", value=text)
                        for top, _bottom, text in overrides
                    ]
                    page_table_count = 0
                else:
                    table_blocks = [block for block in blocks if block.kind == "table"]
                    blocks = [
                        _PageBlock(top=top, kind="paragraph", value=text)
                        for top, bottom, text in overrides
                        if not any(
                            table_top <= (top + bottom) / 2 <= table_bottom
                            for _left, table_top, _right, table_bottom in table_boxes
                        )
                    ]
                    blocks.extend(table_blocks)
                blocks.sort(key=lambda block: block.top)
                used_native_content = False
                used_ocr = True
            elif not blocks:
                ocr_rows = pdf_extraction._ocr_page(pdf_bytes, page_index)
                blocks = [
                    _PageBlock(float(index), "paragraph", " ".join(cell for cell in row if cell))
                    for index, row in enumerate(ocr_rows)
                    if any(row)
                ]
                used_ocr = bool(blocks)

            reader_page = None
            if image_pdf is not None:
                try:
                    reader_page = image_pdf.pages[page_index]
                except Exception:
                    reader_page = None
            image_blocks = _extract_image_blocks(
                page,
                reader_page,
                page_number=page_index + 1,
            )
            blocks.extend(image_blocks)
            blocks.sort(key=lambda block: block.top)

            if not blocks:
                continue
            if converted_page_count:
                document.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

            for block in blocks:
                if block.kind == "table":
                    _append_table(document, block.value)  # type: ignore[arg-type]
                elif block.kind == "image":
                    _append_image(document, block.value)  # type: ignore[arg-type]
                else:
                    _append_paragraph(document, str(block.value))

            converted_page_count += 1
            table_count += page_table_count
            image_count += len(image_blocks)
            if used_ocr:
                ocr_page_count += 1
            elif used_native_content:
                text_page_count += 1
    finally:
        pdf.close()

    if converted_page_count == 0:
        raise PdfToWordConversionError(
            "未识别到可写入 Word 的文字或表格；如果是扫描件，请确认服务器已安装 OCR 引擎并尝试更清晰的 PDF。",
        )

    output = BytesIO()
    document.save(output)
    return PdfToWordResult(
        content=output.getvalue(),
        page_count=page_count,
        table_count=table_count,
        image_count=image_count,
        text_page_count=text_page_count,
        ocr_page_count=ocr_page_count,
        output_file_name=_output_file_name(source_file_name),
    )
