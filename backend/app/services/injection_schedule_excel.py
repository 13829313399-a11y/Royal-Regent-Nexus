from __future__ import annotations

from collections import Counter
from datetime import date, datetime, time
from hashlib import sha256
from io import BytesIO
import json
from pathlib import PurePosixPath
import re
from typing import Any
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile, is_zipfile

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PARSER_VERSION = "injection-daily-schedule-v3"
CORE_SHEETS = ("计划表", "机台完成时间", "机安", "单价")
PLAN_MAX_COLUMN = 49
PLAN_MAX_ROWS = 5_000
MAX_COMPRESSED_BYTES = 32 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
MAX_ARCHIVE_MEMBER_BYTES = 64 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 200
MAX_COMPRESSION_RATIO = 250
MIN_PENDING_SEPARATOR_ROWS = 3
MACHINE_CODE_PATTERN = re.compile(r"^(?:旧|新)\d+$")
TONNAGE_PATTERN = re.compile(r"(\d+(?:\.\d+)?)\s*T", re.IGNORECASE)
FORMULA_ERROR_PATTERN = re.compile(r"^#(?:N/A|REF!|VALUE!|NAME\?|DIV/0!|NUM!|NULL!)$", re.IGNORECASE)


def inspect_xlsx_upload(content: bytes, source_file_name: str, content_type: str = "") -> dict[str, int]:
    if not source_file_name.lower().endswith(".xlsx"):
        raise HTTPException(status_code=415, detail="注塑排产仅支持 .xlsx 工作簿")
    if not content:
        raise HTTPException(status_code=400, detail="上传的工作簿为空")
    if len(content) > MAX_COMPRESSED_BYTES:
        raise HTTPException(
            status_code=413,
            detail={
                "code": "xlsx_too_large",
                "message": "工作簿超过 32 MiB 上传限制",
                "max_bytes": MAX_COMPRESSED_BYTES,
            },
        )
    if content_type and content_type not in {
        XLSX_MIME,
        "application/octet-stream",
        "application/zip",
    }:
        raise HTTPException(status_code=415, detail="上传内容类型不是有效的 Excel 工作簿")
    if not is_zipfile(BytesIO(content)):
        raise HTTPException(status_code=400, detail="工作簿不是有效的 XLSX ZIP 文件")

    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise HTTPException(
                    status_code=413,
                    detail={
                        "code": "xlsx_too_many_members",
                        "message": "工作簿内部文件数量异常",
                    },
                )
            names = {item.filename for item in members}
            if not {"[Content_Types].xml", "xl/workbook.xml"} <= names:
                raise HTTPException(status_code=400, detail="工作簿缺少必要的 XLSX 结构")

            total_uncompressed = 0
            for item in members:
                path = PurePosixPath(item.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise HTTPException(status_code=400, detail="工作簿包含不安全的内部路径")
                if item.flag_bits & 0x1:
                    raise HTTPException(status_code=400, detail="不支持加密的工作簿")
                if item.file_size > MAX_ARCHIVE_MEMBER_BYTES:
                    raise HTTPException(status_code=413, detail="工作簿内部单个文件异常过大")
                if (
                    item.file_size > 1024 * 1024
                    and item.compress_size > 0
                    and item.file_size / item.compress_size > MAX_COMPRESSION_RATIO
                ):
                    raise HTTPException(status_code=413, detail="工作簿压缩比异常，已拒绝解析")
                total_uncompressed += item.file_size
                if total_uncompressed > MAX_UNCOMPRESSED_BYTES:
                    raise HTTPException(status_code=413, detail="工作簿解压后内容超过安全限制")
    except BadZipFile as error:
        raise HTTPException(status_code=400, detail="工作簿 ZIP 结构损坏") from error

    return {
        "compressed_bytes": len(content),
        "uncompressed_bytes": total_uncompressed,
        "member_count": len(members),
    }


def parse_daily_schedule_workbook(
    content: bytes,
    source_file_name: str,
    content_type: str = "",
) -> dict[str, Any]:
    archive_stats = inspect_xlsx_upload(content, source_file_name, content_type)
    external_formula_cells = scan_external_formula_references(content)

    try:
        workbook = load_workbook(
            BytesIO(content),
            data_only=True,
            read_only=True,
            keep_links=False,
        )
    except Exception as error:
        raise HTTPException(status_code=400, detail=f"无法读取 Excel 工作簿：{error}") from error

    try:
        sheet = select_plan_sheet(workbook)
        detected_sheets = [name for name in CORE_SHEETS if name in workbook.sheetnames]
        issues: list[dict[str, Any]] = []
        for sheet_name in CORE_SHEETS:
            if sheet_name not in workbook.sheetnames:
                issues.append(
                    issue(
                        code="missing_core_sheet",
                        severity="error",
                        blocking=True,
                        source_sheet=sheet_name,
                        message=f"缺少核心工作表“{sheet_name}”。",
                    )
                )

        finish_by_machine = parse_machine_finish_sheet(workbook)
        mold_reference_by_code = parse_mold_sheet(workbook)
        price_reference_by_key = parse_price_sheet(workbook)
        plan_rows = list(
            sheet.iter_rows(
                min_row=1,
                max_row=min(sheet.max_row or PLAN_MAX_ROWS, PLAN_MAX_ROWS),
                max_col=PLAN_MAX_COLUMN,
                values_only=True,
            )
        )
        last_data_row = find_last_plan_data_row(plan_rows)
        plan_rows = plan_rows[:last_data_row]

        machine_rows: dict[int, dict[str, Any]] = {}
        for row_index, row in enumerate(plan_rows, start=1):
            if is_machine_row(row):
                parsed_machine = parse_machine_row(
                    row,
                    row_index,
                    finish_by_machine,
                )
                machine_rows[row_index] = parsed_machine
        machines = list(machine_rows.values())
        machine_codes = {item["machine_code"] for item in machines}
        if not machines:
            issues.append(
                issue(
                    code="no_machines",
                    severity="error",
                    blocking=True,
                    source_sheet=sheet.title,
                    message="计划表没有识别到机台主数据。",
                )
            )

        pending_pool_start_row = find_pending_pool_start(plan_rows, machine_rows)
        orders: list[dict[str, Any]] = []
        current_machine_code = ""
        for row_index, row in enumerate(plan_rows, start=1):
            machine = machine_rows.get(row_index)
            if machine is not None:
                current_machine_code = machine["machine_code"]
                continue
            if not is_task_row(row):
                continue

            in_pending_pool = bool(
                pending_pool_start_row and row_index >= pending_pool_start_row
            )
            explicit_machine_code = text(value_at(row, 2))
            if explicit_machine_code not in machine_codes:
                explicit_machine_code = ""
            assigned_machine_code = (
                explicit_machine_code
                if in_pending_pool
                else explicit_machine_code or current_machine_code
            )
            assignment_state = (
                "preassigned"
                if in_pending_pool and assigned_machine_code
                else "unassigned"
                if in_pending_pool
                else "assigned"
                if assigned_machine_code
                else "unassigned"
            )
            parsed_order, row_issues = parse_order_row(
                row=row,
                row_index=row_index,
                source_sheet=sheet.title,
                assigned_machine_code=assigned_machine_code,
                assignment_state=assignment_state,
                in_pending_pool=in_pending_pool,
                price_reference_by_key=price_reference_by_key,
                mold_reference_by_code=mold_reference_by_code,
            )
            orders.append(parsed_order)
            issues.extend(row_issues)

        assign_stable_natural_keys(orders, issues, sheet.title)
        referenced_mold_codes = {
            normalize_key(item["mold_code"]) for item in orders if item["mold_code"]
        }
        molds = [
            reference
            for code, references in mold_reference_by_code.items()
            if code in referenced_mold_codes
            for reference in references[:1]
        ]
        mold_codes = {normalize_key(item["mold_code"]) for item in molds}
        for order in orders:
            code = normalize_key(order["mold_code"])
            if not code or code in mold_codes:
                continue
            molds.append(
                {
                    "mold_code": order["mold_code"],
                    "mold_name": order["product_name"],
                    "machine_class": order["machine_class"],
                    "robot_type": "",
                    "fixture_type": "",
                    "length_mm": None,
                    "width_mm": None,
                    "height_mm": None,
                    "mold_weight_kg": None,
                    "gross_shot_weight_g": order["source_values"].get(
                        "gross_shot_weight_g"
                    ),
                    "cavities": None,
                    "cycle_seconds": None,
                    "required_capabilities": [],
                    "material_rules": [],
                    "quality_status": "incomplete",
                    "source_sheet": sheet.title,
                    "source_row": order["source_row"],
                }
            )
            mold_codes.add(code)

        for entry in external_formula_cells:
            issues.append(
                issue(
                    code="external_formula_reference",
                    severity="warning",
                    blocking=False,
                    source_sheet=entry["source_sheet"],
                    source_row=entry["source_row"],
                    field_name=entry["cell"],
                    raw_value=entry["formula"],
                    message="检测到外部工作簿公式，导入值按当前缓存读取并保留风险标记。",
                )
            )

        business_date = normalize_real_date(value_at(plan_rows[1], 9)) if len(plan_rows) > 1 else ""
        assigned_count = sum(item["assignment_state"] == "assigned" for item in orders)
        preassigned_count = sum(
            item["assignment_state"] == "preassigned" for item in orders
        )
        unassigned_count = sum(
            item["assignment_state"] == "unassigned" for item in orders
        )
        summary = {
            "machine_count": len(machines),
            "old_machine_count": sum(item["workshop"] == "old" for item in machines),
            "new_machine_count": sum(item["workshop"] == "new" for item in machines),
            "order_count": len(orders),
            "assigned_order_count": assigned_count,
            "preassigned_pending_count": preassigned_count,
            "unassigned_order_count": unassigned_count,
            "pending_pool_count": preassigned_count + unassigned_count,
            "mold_count": len(molds),
            "issue_count": len(issues),
            "blocking_issue_count": sum(bool(item["blocking"]) for item in issues),
            "formula_error_count": sum(
                item["code"] == "formula_error" for item in issues
            ),
            "external_formula_risk_count": len(external_formula_cells),
            "business_date": business_date,
            "pending_pool_start_row": pending_pool_start_row,
            "archive": archive_stats,
        }
        return {
            "source_file_name": source_file_name,
            "source_sha256": sha256(content).hexdigest(),
            "parser_version": PARSER_VERSION,
            "detected_sheets": detected_sheets,
            "business_date": business_date,
            "summary": summary,
            "preview": {
                "machines": machines,
                "molds": molds,
                "orders": orders,
            },
            "issues": issues,
        }
    finally:
        workbook.close()


def select_plan_sheet(workbook):
    if "计划表" in workbook.sheetnames:
        return workbook["计划表"]
    for sheet in workbook.worksheets:
        rows = list(
            sheet.iter_rows(
                min_row=1,
                max_row=5,
                max_col=PLAN_MAX_COLUMN,
                values_only=True,
            )
        )
        labels = {text(value) for row in rows for value in row if not is_blank(value)}
        if {"机号", "工模", "名称", "单号", "货号", "欠数", "计划目标"} <= labels:
            return sheet
    raise HTTPException(status_code=422, detail="未找到“计划表”或兼容的排产表头")


def parse_machine_finish_sheet(workbook) -> dict[str, dict[str, Any]]:
    if "机台完成时间" not in workbook.sheetnames:
        return {}
    result: dict[str, dict[str, Any]] = {}
    sheet = workbook["机台完成时间"]
    for row in sheet.iter_rows(min_row=3, max_row=500, max_col=13, values_only=True):
        machine_code = text(value_at(row, 1))
        if not MACHINE_CODE_PATTERN.fullmatch(machine_code):
            continue
        result[machine_code] = {
            "available_at": normalize_real_datetime(value_at(row, 2)),
            "machine_class": text(value_at(row, 3)),
            "process_type": text(value_at(row, 4)),
            "robot_type": text(value_at(row, 5)),
        }
    return result


def parse_mold_sheet(workbook) -> dict[str, list[dict[str, Any]]]:
    if "机安" not in workbook.sheetnames:
        return {}
    result: dict[str, list[dict[str, Any]]] = {}
    sheet = workbook["机安"]
    max_row = min(sheet.max_row or 5_000, 20_000)
    for row_index, row in enumerate(
        sheet.iter_rows(min_row=2, max_row=max_row, max_col=23, values_only=True),
        start=2,
    ):
        mold_code = text(value_at(row, 1))
        normalized = normalize_key(mold_code)
        if not normalized or normalized in {"工模", "工模编号"}:
            continue
        item = {
            "mold_code": mold_code,
            "mold_name": text(value_at(row, 2)),
            "machine_class": text(value_at(row, 7)),
            "robot_type": text(value_at(row, 8)),
            "fixture_type": text(value_at(row, 9)),
            "length_mm": coerce_positive_number(value_at(row, 20)),
            "width_mm": coerce_positive_number(value_at(row, 21)),
            "height_mm": coerce_positive_number(value_at(row, 22)),
            "mold_weight_kg": coerce_positive_number(value_at(row, 23)),
            "gross_shot_weight_g": coerce_positive_number(value_at(row, 13)),
            "cavities": None,
            "cycle_seconds": None,
            "required_capabilities": [],
            "material_rules": [],
            "quality_status": "verified"
            if all(
                coerce_positive_number(value_at(row, column)) is not None
                for column in (20, 21, 22)
            )
            else "incomplete",
            "source_sheet": "机安",
            "source_row": row_index,
        }
        result.setdefault(normalized, []).append(item)
    return result


def parse_price_sheet(workbook) -> dict[tuple[str, str, str], list[dict[str, Any]]]:
    if "单价" not in workbook.sheetnames:
        return {}
    sheet = workbook["单价"]
    first_rows = list(
        sheet.iter_rows(min_row=1, max_row=6, max_col=43, values_only=True)
    )
    header_row, headers = find_header(first_rows, ("工模编号", "工模"))
    if not headers:
        return {}
    mold_col = find_header_column(headers, "工模编号", "工模")
    product_col = find_header_column(headers, "塑胶货号", "货号")
    name_col = find_header_column(headers, "物料名称", "产品", "名称")
    weight_col = find_header_column(headers, "整啤毛重", "整啤净重")
    target_col = find_header_column(headers, "模具日产量", "每日目标", "每班目标")
    model_col = find_header_column(headers, "啤机机型", "机型")
    max_col = max(
        filter(
            None,
            (mold_col, product_col, name_col, weight_col, target_col, model_col),
        ),
        default=1,
    )
    max_row = min(sheet.max_row or 10_000, 50_000)
    result: dict[tuple[str, str, str], list[dict[str, Any]]] = {}
    for row in sheet.iter_rows(
        min_row=header_row + 1,
        max_row=max_row,
        max_col=max_col,
        values_only=True,
    ):
        mold_code = normalize_key(value_at(row, mold_col))
        if not mold_code:
            continue
        product_code = normalize_key(value_at(row, product_col))
        product_name = normalize_key(value_at(row, name_col))
        item = {
            "gross_shot_weight_g": coerce_positive_number(
                value_at(row, weight_col)
            ),
            "daily_target_qty": coerce_positive_number(value_at(row, target_col)),
            "machine_class": text(value_at(row, model_col)),
        }
        for key in (
            (mold_code, product_code, product_name),
            (mold_code, product_code, ""),
            (mold_code, "", ""),
        ):
            result.setdefault(key, []).append(item)
    return result


def parse_machine_row(
    row: tuple[Any, ...],
    row_index: int,
    finish_by_machine: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    machine_code = text(value_at(row, 2))
    finish = finish_by_machine.get(machine_code, {})
    machine_class = text(value_at(row, 7)) or text(finish.get("machine_class"))
    tonnage_match = TONNAGE_PATTERN.search(machine_class)
    return {
        "machine_code": machine_code,
        "machine_name": machine_code,
        "workshop": "old" if machine_code.startswith("旧") else "new",
        "machine_class": machine_class,
        "tonnage_t": float(tonnage_match.group(1)) if tonnage_match else None,
        "process_type": text(value_at(row, 8)) or text(finish.get("process_type")),
        "robot_type": text(value_at(row, 9)) or text(finish.get("robot_type")),
        "fixture_type": "",
        "max_shot_weight_g": None,
        "tie_bar_x_mm": None,
        "tie_bar_y_mm": None,
        "mold_thickness_min_mm": None,
        "mold_thickness_max_mm": None,
        "opening_stroke_mm": None,
        "ejector_stroke_mm": None,
        "status": "available",
        "available_at": text(finish.get("available_at")),
        "capabilities": [],
        "material_rules": [],
        "quality_status": "incomplete",
        "source_sheet": "计划表",
        "source_row": row_index,
        "source_values": {
            "machine_position": json_cell_value(value_at(row, 1)),
            "remark": text(value_at(row, 10)),
        },
    }


def parse_order_row(
    *,
    row: tuple[Any, ...],
    row_index: int,
    source_sheet: str,
    assigned_machine_code: str,
    assignment_state: str,
    in_pending_pool: bool,
    price_reference_by_key: dict[tuple[str, str, str], list[dict[str, Any]]],
    mold_reference_by_code: dict[str, list[dict[str, Any]]],
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    mold_code = text(value_at(row, 7))
    product_name = text(value_at(row, 8))
    order_no = text(value_at(row, 9))
    product_code = text(value_at(row, 10))
    price_matches = price_reference_by_key.get(
        (
            normalize_key(mold_code),
            normalize_key(product_code),
            normalize_key(product_name),
        ),
        price_reference_by_key.get(
            (normalize_key(mold_code), normalize_key(product_code), ""),
            [],
        ),
    )
    unique_price = price_matches[0] if len(price_matches) == 1 else {}
    mold_matches = mold_reference_by_code.get(normalize_key(mold_code), [])
    unique_mold = mold_matches[0] if len(mold_matches) == 1 else {}

    order_qty = max(coerce_number(value_at(row, 12)) or 0, 0)
    produced_qty = max(coerce_number(value_at(row, 13)) or 0, 0)
    shortage = coerce_number(value_at(row, 14))
    outstanding_qty = max(
        shortage if shortage is not None else order_qty - produced_qty,
        0,
    )
    daily_target_qty = (
        coerce_positive_number(value_at(row, 15))
        or unique_price.get("daily_target_qty")
    )
    plan_start_at = (
        "" if in_pending_pool else normalize_real_datetime(value_at(row, 33))
    )
    plan_finish_at = (
        "" if in_pending_pool else normalize_real_datetime(value_at(row, 34))
    )
    due_date = normalize_real_date(value_at(row, 28))
    gross_shot_weight = (
        coerce_positive_number(value_at(row, 21))
        or unique_price.get("gross_shot_weight_g")
        or unique_mold.get("gross_shot_weight_g")
    )
    row_issues: list[dict[str, Any]] = []

    for column in range(1, PLAN_MAX_COLUMN + 1):
        value = value_at(row, column)
        if isinstance(value, str) and FORMULA_ERROR_PATTERN.fullmatch(value.strip()):
            row_issues.append(
                issue(
                    code="formula_error",
                    severity="warning",
                    blocking=False,
                    source_sheet=source_sheet,
                    source_row=row_index,
                    field_name=get_column_letter(column),
                    raw_value=value,
                    message="公式缓存值错误，已保留来源并按缺失资料处理。",
                )
            )

    raw_plan_values = (value_at(row, 33), value_at(row, 34))
    if any(not is_blank(value) for value in raw_plan_values) and not (
        plan_start_at or plan_finish_at
    ):
        row_issues.append(
            issue(
                code=(
                    "pending_pool_formula_ignored"
                    if in_pending_pool
                    else "relative_plan_time"
                ),
                severity="info",
                blocking=False,
                source_sheet=source_sheet,
                source_row=row_index,
                field_name="AG/AH",
                raw_value=" / ".join(stringify_raw(value) for value in raw_plan_values),
                message="检测到纯时间或 1900 日期，正式计划时间已清空。",
            )
        )
    if not due_date:
        row_issues.append(
            issue(
                code="missing_due_date",
                severity="warning",
                blocking=False,
                source_sheet=source_sheet,
                source_row=row_index,
                field_name="AB",
                message="缺少有效交期；订单可导入，但相关排程不能发布。",
            )
        )
    if not order_no and not product_code:
        row_issues.append(
            issue(
                code="missing_order_identity",
                severity="error",
                blocking=True,
                source_sheet=source_sheet,
                source_row=row_index,
                field_name="I/J",
                message="单号和货号同时缺失，无法建立稳定订单标识。",
            )
        )
    if len(price_matches) > 1:
        row_issues.append(
            issue(
                code="ambiguous_price_reference",
                severity="warning",
                blocking=False,
                source_sheet=source_sheet,
                source_row=row_index,
                field_name="G/J",
                message="单价表存在多条候选记录，未自动采用射胶量或日产量。",
            )
        )
    if len(mold_matches) > 1:
        row_issues.append(
            issue(
                code="ambiguous_mold_reference",
                severity="warning",
                blocking=False,
                source_sheet=source_sheet,
                source_row=row_index,
                field_name="G",
                message="机安表存在多条同模号记录，需在主数据中人工确认。",
            )
        )

    machine_class = text(value_at(row, 6)) or text(
        unique_price.get("machine_class")
    )
    quality_status = "verified" if due_date and mold_code and daily_target_qty else "incomplete"
    return (
        {
            "natural_key": "",
            "order_no": order_no,
            "product_code": product_code,
            "product_name": product_name,
            "mold_code": mold_code,
            "color": text(value_at(row, 17)),
            "pigment": text(value_at(row, 18)),
            "material": text(value_at(row, 19)),
            "machine_class": machine_class,
            "order_qty": order_qty,
            "produced_qty": produced_qty,
            "outstanding_qty": outstanding_qty,
            "daily_target_qty": daily_target_qty,
            "delivery_due_date": due_date,
            "priority_flag": text(value_at(row, 5)),
            "status": "completed" if outstanding_qty <= 0 else "open",
            "assigned_machine_code": assigned_machine_code,
            "assignment_state": assignment_state,
            "plan_start_at": plan_start_at,
            "plan_finish_at": plan_finish_at,
            "quality_status": quality_status,
            "source_sheet": source_sheet,
            "source_row": row_index,
            "source_values": {
                "quantity_set": coerce_number(value_at(row, 11)),
                "gross_shot_weight_g": gross_shot_weight,
                "required_material_kg": coerce_positive_number(value_at(row, 22)),
                "warehouse_due_at": normalize_real_datetime(value_at(row, 36)),
                "delivery_gap_days": coerce_number(value_at(row, 37)),
                "remark": text(value_at(row, 49)),
                "raw_cells": {
                    str(column): json_cell_value(value_at(row, column))
                    for column in range(1, PLAN_MAX_COLUMN + 1)
                    if not is_blank(value_at(row, column))
                },
            },
        },
        row_issues,
    )


def assign_stable_natural_keys(
    orders: list[dict[str, Any]],
    issues: list[dict[str, Any]],
    source_sheet: str,
) -> None:
    business_bases = [
        "|".join(
            normalize_key(order[field])
            for field in ("order_no", "product_code", "mold_code")
        )
        for order in orders
    ]
    identity_bases = [
        "|".join(
            (
                business_base,
                normalize_key(order["color"]),
                normalize_key(order["pigment"]),
                normalized_number_key(order["order_qty"]),
            )
        )
        for order, business_base in zip(orders, business_bases, strict=True)
    ]
    business_totals = Counter(business_bases)
    identity_totals = Counter(identity_bases)
    for order, business_base, identity_base in zip(
        orders,
        business_bases,
        identity_bases,
        strict=True,
    ):
        key_source = (
            identity_base
            if identity_totals[identity_base] == 1
            else f"{identity_base}|source-row:{order['source_row']}"
        )
        order["natural_key"] = sha256(key_source.encode("utf-8")).hexdigest()[:40]
        if business_totals[business_base] > 1:
            issues.append(
                issue(
                    code="duplicate_order_identity",
                    severity="warning",
                    blocking=False,
                    source_sheet=source_sheet,
                    source_row=order["source_row"],
                    field_name="G/I/J/Q/R/K",
                    raw_value=business_base,
                    message="同一订单基础业务键重复，已按颜色、色粉及订单数量稳定区分分项。",
                )
            )
        if identity_totals[identity_base] > 1:
            issues.append(
                issue(
                    code="ambiguous_order_identity",
                    severity="error",
                    blocking=True,
                    source_sheet=source_sheet,
                    source_row=order["source_row"],
                    field_name="G/I/J/Q/R/K",
                    raw_value=identity_base,
                    message="订单基础键、颜色、色粉及数量仍完全重复，无法安全匹配历史订单。",
                )
            )


def normalized_number_key(value: Any) -> str:
    number = coerce_number(value)
    if number is None:
        return ""
    return f"{number:.6f}".rstrip("0").rstrip(".")


def find_pending_pool_start(
    rows: list[tuple[Any, ...]],
    machine_rows: dict[int, dict[str, Any]],
) -> int:
    if not machine_rows:
        return 0
    last_machine_row = max(machine_rows)
    blank_run = 0
    for row_index in range(last_machine_row + 1, len(rows) + 1):
        row = rows[row_index - 1]
        if is_task_row(row):
            if blank_run >= MIN_PENDING_SEPARATOR_ROWS:
                return row_index
            blank_run = 0
        elif is_machine_row(row):
            blank_run = 0
        elif is_empty_plan_row(row):
            blank_run += 1
        else:
            blank_run = 0
    return 0


def find_last_plan_data_row(rows: list[tuple[Any, ...]]) -> int:
    last_row = 0
    for row_index, row in enumerate(rows, start=1):
        if any(not is_blank(value_at(row, column)) for column in range(1, 16)):
            last_row = row_index
    return last_row


def is_machine_row(row: tuple[Any, ...]) -> bool:
    machine_code = text(value_at(row, 2))
    quantity_values = [value_at(row, column) for column in range(11, 16)]
    return bool(
        MACHINE_CODE_PATTERN.fullmatch(machine_code)
        and text(value_at(row, 7))
        and text(value_at(row, 8))
        and text(value_at(row, 9))
        and all(coerce_number(value) is None for value in quantity_values)
    )


def is_task_row(row: tuple[Any, ...]) -> bool:
    core_values = [value_at(row, column) for column in range(7, 11)]
    quantity_values = [value_at(row, column) for column in range(11, 16)]
    return sum(not is_blank(value) for value in core_values) >= 3 and any(
        coerce_number(value) is not None for value in quantity_values
    )


def is_empty_plan_row(row: tuple[Any, ...]) -> bool:
    return all(is_blank(value_at(row, column)) for column in range(1, 16))


def scan_external_formula_references(content: bytes) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    cell_tag = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c"
    formula_tag = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f"
    try:
        with ZipFile(BytesIO(content)) as archive:
            plan_sheet_path = resolve_plan_sheet_archive_path(archive)
            sheet_names = (
                [plan_sheet_path]
                if plan_sheet_path
                else sorted(
                    name
                    for name in archive.namelist()
                    if name.startswith("xl/worksheets/") and name.endswith(".xml")
                )[:1]
            )
            for sheet_name in sheet_names:
                with archive.open(sheet_name) as stream:
                    for _, element in ElementTree.iterparse(stream, events=("end",)):
                        if element.tag != cell_tag:
                            continue
                        formula = element.find(formula_tag)
                        formula_text = formula.text if formula is not None else ""
                        if formula_text and (
                            "[" in formula_text or ".xls" in formula_text.lower()
                        ):
                            coordinate = element.attrib.get("r", "")
                            row_match = re.search(r"(\d+)$", coordinate)
                            results.append(
                                {
                                    "source_sheet": sheet_name,
                                    "source_row": int(row_match.group(1))
                                    if row_match
                                    else 0,
                                    "cell": coordinate,
                                    "formula": formula_text[:500],
                                }
                            )
                            if len(results) >= 100:
                                return results
                        element.clear()
    except (BadZipFile, ElementTree.ParseError):
        return []
    return results


def resolve_plan_sheet_archive_path(archive: ZipFile) -> str:
    try:
        workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        relationships_root = ElementTree.fromstring(
            archive.read("xl/_rels/workbook.xml.rels")
        )
    except (KeyError, ElementTree.ParseError):
        return ""

    relation_by_id = {
        item.attrib.get("Id", ""): item.attrib.get("Target", "")
        for item in relationships_root
    }
    relationship_id_key = (
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    )
    sheets = workbook_root.findall(
        ".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet"
    )
    selected = next(
        (item for item in sheets if item.attrib.get("name") == "计划表"),
        sheets[0] if sheets else None,
    )
    if selected is None:
        return ""
    target = relation_by_id.get(selected.attrib.get(relationship_id_key, ""), "")
    if not target:
        return ""
    normalized = str(PurePosixPath("xl") / target).replace("xl/../", "")
    return normalized if normalized in archive.namelist() else ""


def find_header(
    rows: list[tuple[Any, ...]],
    candidates: tuple[str, ...],
) -> tuple[int, tuple[Any, ...]]:
    for row_index, row in enumerate(rows, start=1):
        normalized = {text(value) for value in row}
        if any(candidate in normalized for candidate in candidates):
            return row_index, row
    return 0, ()


def find_header_column(headers: tuple[Any, ...], *labels: str) -> int:
    values = [text(value) for value in headers]
    for label in labels:
        if label in values:
            return values.index(label) + 1
    return 0


def issue(
    *,
    code: str,
    severity: str,
    blocking: bool,
    source_sheet: str = "",
    source_row: int = 0,
    field_name: str = "",
    raw_value: str = "",
    message: str,
) -> dict[str, Any]:
    return {
        "source_sheet": source_sheet,
        "source_row": source_row,
        "severity": severity,
        "code": code,
        "field_name": field_name,
        "blocking": blocking,
        "raw_value": raw_value,
        "message": message,
    }


def value_at(row: tuple[Any, ...], one_based_column: int) -> Any:
    if one_based_column <= 0 or one_based_column > len(row):
        return None
    return row[one_based_column - 1]


def coerce_number(value: Any) -> float | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def coerce_positive_number(value: Any) -> float | None:
    parsed = coerce_number(value)
    return parsed if parsed is not None and parsed > 0 else None


def normalize_real_date(value: Any) -> str:
    normalized = normalize_datetime(value)
    if normalized is None or normalized.year < 2020:
        return ""
    return normalized.date().isoformat()


def normalize_real_datetime(value: Any) -> str:
    normalized = normalize_datetime(value)
    if normalized is None or normalized.year < 2020:
        return ""
    return normalized.strftime("%Y-%m-%d %H:%M:%S")


def normalize_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    return None


def normalize_key(value: Any) -> str:
    return re.sub(r"\s+", "", text(value)).upper()


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def stringify_raw(value: Any) -> str:
    if isinstance(value, datetime):
        return value.time().isoformat() if value.year < 2020 else value.isoformat()
    if isinstance(value, date):
        return "" if value.year < 2020 else value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    return text(value)


def json_cell_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.time().isoformat() if value.year < 2020 else value.isoformat()
    if isinstance(value, date):
        return "" if value.year < 2020 else value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def canonical_json(value: Any) -> str:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
