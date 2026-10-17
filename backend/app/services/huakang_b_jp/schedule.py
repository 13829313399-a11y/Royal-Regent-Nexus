# Ported from huakang-jp-entry d82c172; server-only adapter, no desktop startup.
# -*- coding: utf-8 -*-
"""排期参照：从导出的排货报表建立索引，用于查重、客名→国家带出、新货号预警。"""
import os
import sys
from collections import Counter, defaultdict

import openpyxl


def _base_dir():
    # PyInstaller 打包后数据要落在 exe 旁边（可写），而不是临时解压目录
    if getattr(sys, 'frozen', False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


DATA_DIR = ''  # No desktop persistence in the server adapter.
REF_FILE = os.path.join(DATA_DIR, 'schedule_ref.xlsx')
REF_META = os.path.join(DATA_DIR, 'schedule_ref.json')

# 报表中表头在第 3 行、数据从第 4 行开始的工作表（列布局一致）
REF_SHEETS = ['问单', '华康B', '华康B已验货', '车衣部', '车衣部已验货']

COL_客名, COL_客PO, COL_货号, COL_数量, COL_国家, COL_版本 = 1, 2, 5, 10, 13, 9  # 0-based
COL_子货号, COL_报关货号, COL_产品名称, COL_装箱数 = 6, 7, 8, 12


class ScheduleRef:
    def __init__(self):
        self.loaded = False
        self.name = ''
        self.uploaded_at = ''
        self.row_count = 0
        self.known_items = set()          # 货号集合（判断 T 列能否带出资料）
        self.country_by_customer = {}     # 客名 → 最常见国家
        self.po_item_keys = set()         # (客PO, 货号)
        self.cust_item_qty_keys = set()   # (客名, 货号, 数量)
        self.max_seq = 0                  # 订单编号最大流水号（如 0126-21614 → 21614）
        self.version_by_item = {}         # 货号 → 唯一历史版本（多版本不带出）
        self.item_detail = {}             # 报关货号 → 货号/子货号/产品名称/装箱数（填 T 列用）

    def load(self, path=REF_FILE, name='', uploaded_at=''):
        import re
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        if not any(s in wb.sheetnames for s in REF_SHEETS):
            wb.close()
            raise ValueError('文件中没有「问单 / 华康B」等工作表——请上传排货报表，而不是金山订单表的导出')
        self.known_items.clear()
        self.po_item_keys.clear()
        self.cust_item_qty_keys.clear()
        self.item_detail.clear()
        self.max_seq = 0
        country_votes = defaultdict(Counter)
        version_votes = defaultdict(Counter)
        n = 0
        for sheet in REF_SHEETS:
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            if ws.max_row and ws.max_row > 200000:
                wb.close()
                raise ValueError("排期参照行数或格式范围过大，请先整理导出")
            for index, cells in enumerate(ws.iter_rows(min_row=4, max_col=14), start=4):
                if index > 200000:
                    wb.close()
                    raise ValueError("排期参照超过 20 万行，请先整理")
                row = [c.value for c in cells]
                # Preserve zero-padded identifiers from the original workbook.
                for column in (0, 2, 5, 6, 7):
                    cell = cells[column]
                    value, fmt = cell.value, cell.number_format
                    if isinstance(value, (int, float)) and not isinstance(value, bool) and re.fullmatch(r"0+", fmt or ""):
                        if float(value).is_integer():
                            row[column] = str(int(value)).zfill(len(fmt))
                if not row or row[0] in (None, '', '汇总'):
                    continue
                m = re.match(r'^\d{4}-(\d{4,5})$', str(row[0]).strip())
                if m:
                    self.max_seq = max(self.max_seq, int(m.group(1)))
                item = str(row[COL_货号] or '').strip()
                cust = str(row[COL_客名] or '').strip()
                po = str(row[COL_客PO] or '').strip()
                qty = row[COL_数量]
                country = str(row[COL_国家] or '').strip()
                if not item or item in ('#REF!', '0'):
                    continue
                n += 1
                self.known_items.add(item)
                if po:
                    self.po_item_keys.add((po, item))
                if cust:
                    try:
                        q = int(float(qty)) if qty not in (None, '') else 0
                    except (ValueError, TypeError):
                        q = 0
                    self.cust_item_qty_keys.add((cust, item, q))
                if cust and country:
                    country_votes[cust][country] += 1
                version = str(row[COL_版本] or '').strip()
                if version:
                    version_votes[item][version] += 1
                # 货号资料（填 T 列）：后扫到的工作表（生产单）覆盖前面的（问单）
                customs = str(row[COL_报关货号] or '').strip()
                detail = {
                    '货号': item,
                    '子货号': str(row[COL_子货号] or '').strip(),
                    '产品名称': str(row[COL_产品名称] or '').strip(),
                    '装箱数': str(row[COL_装箱数] or '').strip(),
                }
                self.item_detail[customs or item] = detail
        wb.close()
        # 国家带出：该客名历史国家高度集中（≥80%）才带出，像 JPEU 这种多国/常留空的不带
        self.country_by_customer = {}
        for c, cnt in country_votes.items():
            top, topn = cnt.most_common(1)[0]
            if topn / sum(cnt.values()) >= 0.8:
                self.country_by_customer[c] = top
        # 版本建议：该货号历史只有唯一版本时才带出，多个版本宁可不带
        self.version_by_item = {}
        for item, cnt in version_votes.items():
            if len(cnt) == 1:
                self.version_by_item[item] = cnt.most_common(1)[0][0]
        self.loaded = True
        self.name = name
        self.uploaded_at = uploaded_at
        self.row_count = n

    def check_row(self, customer, customer_po, baby_item, qty):
        """返回 (exists, known_item, country_suggest, version_suggest)。"""
        baby_item = str(baby_item or '').strip()
        customer_po = str(customer_po or '').strip()
        customer = str(customer or '').strip()
        known = baby_item in self.known_items if baby_item else False
        exists = False
        if customer_po:
            # 有客PO 的行只用 客PO+货号 判重，不用兜底（防止误伤同货号同数量的新单）
            exists = (customer_po, baby_item) in self.po_item_keys
        else:
            try:
                q = int(float(qty)) if qty not in (None, '') else 0
            except (ValueError, TypeError):
                q = 0
            if customer and (customer, baby_item, q) in self.cust_item_qty_keys:
                exists = True
        suggest = self.country_by_customer.get(customer, '')
        version = self.version_by_item.get(baby_item, '')
        return exists, known, suggest, version

    def get_item_detail(self, baby_item):
        """报关货号 → {货号, 子货号, 产品名称, 装箱数}，没有返回 None。"""
        return self.item_detail.get(str(baby_item or '').strip())

    def info(self):
        from datetime import date
        prefix = '01' + str(date.today().year)[2:]
        return {'loaded': self.loaded, 'name': self.name,
                'uploaded_at': self.uploaded_at, 'row_count': self.row_count,
                'known_items': len(self.known_items),
                'customers': len(self.country_by_customer),
                'next_seq': self.max_seq + 1 if self.loaded else 0,
                'seq_prefix': prefix}
