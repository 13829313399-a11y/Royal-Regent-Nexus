"""Bounded, stateless parsing; shared references and numbering live in Nexus."""
from __future__ import annotations
from copy import deepcopy
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
import tempfile
from uuid import uuid4
from zipfile import ZipFile, BadZipFile
from zoneinfo import ZoneInfo

from fastapi import HTTPException
import openpyxl
from sqlalchemy import select, func
from sqlalchemy.orm import Session

from app.models.customer_order_ledger import OrderLedgerSource as Source
from app.models.huakang_b_jp import JPReference, JPReferenceEvent, JPCounter, JPExport, JPExportIdentity
from app.services.transaction_lock import lock_transaction
from . import parser
from .annotate import annotate
from .schedule import ScheduleRef
from .prices import PriceTable

FACTORY = "huakang-b"
CUSTOMER = "jp-huakang-b"
KINDS = {"schedule", "price", "confirmed"}
MAX_FILE = 20 * 1024 * 1024
MAX_TOTAL = 50 * 1024 * 1024
MAX_ROWS = 250
COLUMNS = ["订单编号","客名","客PO","来单日期","合同","货号T","客户T","集团分厂T","子货号T",
 "产品名称T","报关货号","版本","来单数量","利宝","装箱数T","国家","日期码","行Q期D","行Q期备注",
 "截数期","货价(HKD)T","产值(HKD)T","总产值(HKD)","出厂价(HKD)T","抽办货期","客Q货期","箱唛",
 "订单备注","生产备注","备料备注","MO/转单","来料采购T","ERP系统","化验报告","备注","来料采购",
 "接单类型","做货厂区","填写检查","业务跟单","订单创建人","创建时间T","创建时间","更新时间"]


def digest(value):
    return sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()


def archive(content, *, xlsx=False):
    try:
        z = ZipFile(BytesIO(content))
        entries = z.infolist()
        if len(entries) > (3000 if xlsx else 200):
            raise ValueError("压缩包文件数量过多")
        if any(e.flag_bits & 1 for e in entries):
            raise ValueError("不支持加密文件")
        if sum(e.file_size for e in entries) > (200 * 1024**2 if xlsx else MAX_TOTAL):
            raise ValueError("压缩包展开后过大")
        if any(e.file_size > (80 * 1024**2 if xlsx else MAX_FILE) or
               (e.file_size > 1024**2 and e.file_size / max(1, e.compress_size) > 500) for e in entries):
            raise ValueError("文件压缩比例或大小异常")
        if xlsx and "[Content_Types].xml" not in z.namelist():
            raise ValueError("不是有效的 XLSX 文件")
        return z
    except BadZipFile as exc:
        raise ValueError("文件损坏或不是支持的压缩包") from exc


def _jp_validate_reference(kind, content, options=None):
    if kind not in KINDS:
        raise ValueError("不支持的参照类型")
    with archive(content, xlsx=True):
        pass
    with tempfile.TemporaryDirectory(prefix="jp-reference-") as folder:
        p = Path(folder) / "reference.xlsx"
        p.write_bytes(content)
        if kind == "schedule":
            obj = ScheduleRef()
            obj.load(p)
            return {"rows": obj.row_count, "max_seq": obj.max_seq}
        obj = PriceTable()
        if kind == "price":
            obj.import_xlsx(p, source_mtime=0)
            for record in obj.records:
                number(record["price_hk"], optional=True)
            return {"rows": len(obj.records)}
        filled, skipped = obj.import_confirmed_xlsx(p, source_mtime=0, default_currency=(options or {}).get("currency"))
        if skipped:
            raise ValueError("确认表存在无效确认价")
        for record in obj.confirmed.values():
            number(record["price"])
        return {"rows": filled}


def empty_confirmed():
    w = openpyxl.Workbook()
    w.active.append(["产品货号", "确认价格"])
    out = BytesIO()
    w.save(out)
    return out.getvalue()


def reference_state(db):
    out = {}
    for ref in db.scalars(select(JPReference).where(JPReference.factory_id == FACTORY)):
        source = db.get(Source, ref.source_id)
        if not source or source.factory_id != FACTORY or source.customer_code != CUSTOMER:
            raise HTTPException(409, "参照关联异常，请联系管理员")
        out[ref.kind] = {"id": source.id, "name": source.file_name, "sha256": source.sha256,
                         "revision": ref.revision, "actor": ref.actor, "updated_at": ref.updated_at, "options": ref.options}
    return out


def snapshots(db):
    state = reference_state(db)
    if not {"schedule", "price"}.issubset(state):
        raise HTTPException(400, "请先上传华康B排期参照和货价表")
    return [(kind, item, bytes(db.get(Source, item["id"]).content)) for kind, item in sorted(state.items())]


def save_reference(db, kind, name, content, expected, actor, options=None):
    options = options or {}
    lock_transaction(db, "jp-intake", FACTORY)
    current = db.get(JPReference, (FACTORY, kind))
    if (current.revision if current else 0) != expected:
        raise HTTPException(409, "参照已更新，请刷新后重新上传")
    checksum = sha256(content).hexdigest()
    source = db.scalar(select(Source).where(Source.factory_id == FACTORY, Source.customer_code == CUSTOMER,
        Source.kind == "jp-ref-" + kind, Source.sha256 == checksum))
    if current and source and current.source_id == source.id and current.options == options:
        return reference_state(db)
    if not source:
        source = Source(id=uuid4().hex, factory_id=FACTORY, customer_code=CUSTOMER,
            kind="jp-ref-" + kind, sha256=checksum, file_name=name, content=content)
        db.add(source)
        db.flush()
    now = datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(timespec="seconds")
    revision = expected + 1
    if current:
        current.source_id, current.revision, current.actor, current.updated_at = source.id, revision, actor, now
        current.options = options
    else:
        db.add(JPReference(factory_id=FACTORY, kind=kind, source_id=source.id,
            revision=revision, actor=actor, updated_at=now, options=options))
    db.add(JPReferenceEvent(id=uuid4().hex, factory_id=FACTORY, kind=kind, revision=revision,
        source_id=source.id, actor=actor, created_at=now, options=options))
    db.flush()
    return reference_state(db)


def expand(files):
    out = []
    for name, content in files:
        if name.lower().endswith(".zip"):
            with archive(content) as z:
                for entry in z.infolist():
                    if entry.is_dir() or entry.filename.startswith("__MACOSX/"):
                        continue
                    if entry.filename.lower().endswith(".eml") and not Path(entry.filename).name.startswith("._"):
                        out.append((name + "::" + entry.filename, z.read(entry)))
        elif name.lower().endswith(".eml"):
            out.append((name, content))
        else:
            raise ValueError("邮件只支持 EML 或包含 EML 的 ZIP")
    if len(out) > 100 or sum(len(c) for _, c in out) > MAX_TOTAL:
        raise ValueError("每批最多 100 封邮件，展开后合计不超过 50MB")
    return out


def number(value, *, integer=False, optional=False):
    s = ("" if value is None else str(value)).replace(",", "").strip()
    if optional and not s:
        return ""
    try:
        if not re.fullmatch(r"\d+(?:\.\d+)?", s):
            raise ValueError()
        n = Decimal(s)
        if not n.is_finite() or n < 0 or (integer and n == 0) or n >= Decimal("100000000000") or n.as_tuple().exponent < -6:
            raise ValueError()
        if integer and n != n.to_integral():
            raise ValueError()
        return format(n, "f")
    except (ValueError, InvalidOperation):
        raise ValueError("数量/装箱数须为正整数，价格不能为负数，最多 6 位小数")


def day(value):
    s = str(value or "").strip().replace("/", "-")
    if not s:
        return ""
    bits = s.split("-")
    try:
        d = date(*(int(b) for b in bits))
        return d.strftime("%Y/%m/%d")
    except (ValueError, TypeError):
        raise ValueError("日期须为年/月/日")


def _jp_preview(files, html, subject, received_date, references, quantity_selections, actor_id, merchandiser="", identity_overrides=()):
    received = date.fromisoformat(received_date)
    choices = {(x["source_id"], x["table_index"]): x["column_index"] for x in quantity_selections}
    if len(choices) != len(quantity_selections):
        raise ValueError("数量列选择重复")
    with tempfile.TemporaryDirectory(prefix="jp-preview-") as folder:
        ref, prices = ScheduleRef(), PriceTable()
        for kind, metadata, content in references:
            with archive(content, xlsx=True):
                pass
            p = Path(folder) / (kind + ".xlsx")
            p.write_bytes(content)
            if kind == "schedule":
                ref.load(p, metadata["name"])
            elif kind == "price":
                prices.import_xlsx(p, metadata["name"], source_mtime=0)
            else:
                prices.import_confirmed_xlsx(p, metadata["name"], source_mtime=0,
                    default_currency=(metadata.get("options") or {}).get("currency"))
        entries, needs = [], []
        expanded = expand(files)
        if html.strip():
            expanded.append(("粘贴正文", html.encode()))
        if not expanded:
            raise ValueError("请上传邮件或粘贴邮件正文")
        for name, content in expanded:
            source_id = digest([name, sha256(content).hexdigest()])
            selected = {i: c for (sid, i), c in choices.items() if sid == source_id}
            try:
                if name == "粘贴正文":
                    body, title, stamp = html, subject, received
                else:
                    rich, plain, title, stamp = parser.parse_eml(content)
                    body = rich or plain or ""
                    stamp = stamp or received
                result = parser.parse_order_email(body, title, stamp, name, quantity_columns=selected)
            except parser.QuantityColumnRequired as exc:
                table = exc.table[:5]
                needs.append({"source_id": source_id, "source": name, "table_index": exc.table_index,
                    "columns": [{"index": i, "sample": " / ".join(str(row[i]) for row in table if i < len(row))[:240]}
                                for i in range(min(100, max(map(len, table), default=0)))]})
                result = {"orders": [], "warnings": [str(exc)], "meta": {"source": name, "subject": ""}}
                stamp = received
            except Exception as exc:
                result = {"orders": [], "warnings": ["邮件无法解析：" + str(exc)[:200]], "meta": {"source": name, "subject": ""}}
                stamp = received
            result["source_id"] = source_id
            result["meta"].setdefault("source", name)
            # Normalize timestamps before comparing threads; never compare naive
            # and aware datetimes or choose arbitrary tied, different content.
            thread = result["meta"].get("thread_date")
            try:
                stamp = datetime.fromisoformat(thread) if thread else stamp
                if not isinstance(stamp, datetime):
                    stamp = datetime.combine(stamp, datetime.min.time(), tzinfo=ZoneInfo("Asia/Shanghai"))
                stamp = stamp.replace(tzinfo=ZoneInfo("Asia/Shanghai")) if not stamp.tzinfo else stamp
                rank = stamp.astimezone(timezone.utc).isoformat()
            except Exception:
                rank = ""
            entries.append((result, rank, sha256(content).hexdigest()))
        groups = {}
        seen_content = set()
        for i, (res, rank, checksum) in enumerate(entries):
            if checksum in seen_content:
                res['orders'] = []
                res['skipped'] = True
                res['warnings'].append('完全相同的重复邮件已跳过')
                continue
            seen_content.add(checksum)
            key = parser.normalize_subject(res["meta"].get("subject"))
            if key and res["orders"]:
                identities = frozenset((str(o.get('customer_po') or '').strip(), str(o.get('long_code') or '').strip())
                                       for o in res['orders'])
                # A generic subject is not order identity. Keep disjoint POs;
                # without explicit POs, ambiguous revisions require human review.
                groups.setdefault((key, identities), []).append(i)
        for (_, identities), indexes in groups.items():
            if len(indexes) > 1 and any(not po for po, _ in identities):
                for i in indexes:
                    entries[i][0]['orders'] = []
                    entries[i][0]['warnings'].append('同主题邮件缺少明确客PO，不能自动认定修订；请分开核对并只上传需要录入的邮件')
                continue
            latest_rank = max(entries[i][1] for i in indexes)
            newest = [i for i in indexes if entries[i][1] == latest_rank]
            if len({entries[i][2] for i in newest}) > 1:
                for i in indexes:
                    entries[i][0]["orders"] = []
                    entries[i][0]["warnings"].append("同主题同时间存在不同版本，请分开上传并核对")
                continue
            for i in indexes:
                if i != newest[0]:
                    entries[i][0]["orders"] = []
                    entries[i][0]["skipped"] = True
                    entries[i][0]["warnings"].append("同线程较旧或重复邮件已跳过")
        rows, warnings, failed = [], [], []
        for res, _, checksum in entries:
            if not res["orders"] and not res.get("skipped"):
                failed.append(res["meta"]["source"])
            annotate(res, ref, prices)
            warnings.extend(res["warnings"])
            for i, order in enumerate(res["orders"]):
                row = dict(order, id=digest([res["source_id"], i]), source=res["meta"]["source"], source_hash=checksum,
                    order_date=res["meta"].get("order_date") or received.strftime("%Y/%m/%d"),
                    note="，".join(filter(None, [res["meta"].get("ddd"), order.get("item_note"), res["meta"].get("global_note")])))
                if any(len(str(value)) > 1000 for value in row.values()):
                    raise ValueError("问单字段过长，请整理邮件或拆分后再上传")
                effective = dict(row)
                for change in identity_overrides:
                    if change["id"] == row["id"] and change["field"] in ("customer", "customer_po", "qty"):
                        effective[change["field"]] = change["value"].strip()
                row["reference_duplicate"] = ref.check_row(effective.get("customer"), effective.get("customer_po"),
                    row["baby_item"], effective.get("qty"))[0]
                rows.append(row)
        if len(rows) > MAX_ROWS:
            raise ValueError("每批最多 250 条问单明细，请拆分批次")
        fingerprint = digest({"files": [(n, sha256(c).hexdigest()) for n, c in files],
            "html": html, "subject": subject, "received": received_date, "choices": quantity_selections,
            "references": [(k, m["id"], m["revision"]) for k, m, _ in references], "actor": actor_id,
            "merchandiser": merchandiser})
        return {"factory_id": FACTORY, "fingerprint": fingerprint, "rows": rows,
            "inputs": {"received_date": received_date, "subject": subject, "quantity_selections": quantity_selections,
                       "merchandiser": merchandiser},
            "warnings": warnings, "failed_sources": failed, "quantity_selections": needs,
            "reference_versions": {k: m["revision"] for k, m, _ in references}, "max_seq": ref.max_seq}


def resolve_rows(preview, selected, overrides):
    by_id = {r["id"]: deepcopy(r) for r in preview["rows"]}
    if not selected or len(set(selected)) != len(selected) or any(i not in by_id for i in selected):
        raise HTTPException(400, "请选择当前预览中的有效明细，不能重复选择")
    if preview["failed_sources"] or preview["quantity_selections"]:
        raise HTTPException(400, "存在未识别邮件或未确认数量列，请先处理或移除该邮件再预览")
    seen = set()
    for change in overrides:
        key = (change["id"], change["field"])
        if change["id"] not in selected or key in seen:
            raise HTTPException(400, "人工更正重复或不属于所选明细")
        seen.add(key)
        by_id[change["id"]][change["field"]] = change["value"].strip()
    rows = []
    identities = set()
    for id in selected:
        row = by_id[id]
        source = next(r for r in preview['rows'] if r['id'] == id)
        context_changed = any(str(row.get(f) or '').strip() != str(source.get(f) or '').strip()
                              for f in ('customer', 'carton_t'))
        if context_changed:
            if any((id, field) not in seen for field in ('price_t', 'factory_price_t')):
                raise HTTPException(400, '更改客名或装箱数后，必须重新填写核对货价与出厂价，不能沿用旧价格')
            row['price_note'] = '客名或装箱数已更正，货价与出厂价由人工重新核对'
            row['manual_price_review'] = True
        if row.get("exists") or row.get("reference_duplicate"):
            raise HTTPException(409, "排期已有该明细，请在原问单上转确认单，不重复导出")
        try:
            row["qty"] = number(row["qty"], integer=True)
            row["carton_t"] = number(row.get("carton_t"), integer=True)
            row["price_t"] = number(row.get("price_t"))
            row["factory_price_t"] = number(row.get("factory_price_t"), optional=True)
            row["ship_date"], row["cutoff_date"] = day(row.get("ship_date")), day(row.get("cutoff_date"))
            row["order_date"] = day(row["order_date"])
            if day(row["ship_date"]) > day(row["cutoff_date"]):
                raise ValueError("行Q期不能晚于截数期")
        except ValueError as exc:
            raise HTTPException(400, f"{row.get('baby_item')}：{exc}") from exc
        if not row.get("baby_item") or not row.get("customer") or not row["cutoff_date"] or not row["ship_date"]:
            raise HTTPException(400, "货号、客名、行Q期及截数期不能为空，请核对后填写")
        identity = (row["customer"].casefold(), row.get("customer_po") or "", row["long_code"], row["qty"])
        if identity in identities:
            raise HTTPException(409, "所选问单包含重复明细，请核对后只保留一条")
        identities.add(identity)
        rows.append(row)
    return rows


def workbook(rows, first, prefix, actor, merchandiser=""):
    w = openpyxl.Workbook()
    s = w.active
    s.title = "录入新订单"
    s.append(COLUMNS)
    for offset, row in enumerate(rows):
        values = [f"{prefix}-{first + offset:04d}", row["customer"], row.get("customer_po"), row["order_date"],
            row.get("sales_order", ""), row.get("item_t", ""), "JP", "华康B", row.get("sub_item_t", ""),
            row.get("product_t", ""), row["baby_item"], row.get("pkg", ""), int(Decimal(row["qty"])), "",
            int(Decimal(row["carton_t"])), row.get("country", ""), "", row["ship_date"], "", row["cutoff_date"],
            Decimal(row["price_t"]), "", "", Decimal(row["factory_price_t"]) if row["factory_price_t"] else "",
            "", "", "", row.get("note", ""), "", "", "", "", "", "", "", "", "问单", "华康B", "已核对", merchandiser, actor, "", "", ""]
        for index in (3, 17, 19):
            values[index] = datetime.strptime(values[index], "%Y/%m/%d")
        s.append(values)
        for cell in s[s.max_row]:
            if isinstance(cell.value, str):
                # Text from an email is never an Excel formula or external link.
                cell.data_type = "s"
        for index in (4, 18, 20):
            s.cell(s.max_row, index).number_format = "yyyy/mm/dd"
    b = BytesIO()
    w.save(b)
    return b.getvalue()


def save_export(db, preview, selected, overrides, references, files, html, reason, actor, creator_name=""):
    rows = resolve_rows(preview, selected, overrides)
    key = digest([preview["fingerprint"], selected, overrides])
    lock_transaction(db, "jp-intake", FACTORY)
    current = reference_state(db)
    if {k: v["revision"] for k, v in current.items()} != preview["reference_versions"]:
        raise HTTPException(409, "排期或货价参照已更新，请重新解析核对")
    existing = db.scalar(select(JPExport).where(JPExport.factory_id == FACTORY, JPExport.request_hash == key))
    if existing:
        return existing
    identities = [digest(["po", row["customer"].casefold(), row["customer_po"], row["long_code"]])
        if row.get("customer_po") else digest(["no-po", row["source_hash"], row["customer"].casefold(), row["long_code"]])
        for row in rows]
    if len(set(identities)) != len(identities):
        raise HTTPException(409, "所选明细包含重复客PO与货号")
    if db.scalar(select(JPExportIdentity.export_id).where(JPExportIdentity.factory_id == FACTORY,
            JPExportIdentity.identity_key.in_(identities)).limit(1)):
        raise HTTPException(409, "同客PO和长货号已有导出记录；请下载原记录，改单应在原问单上处理，勿重复占号")
    counter = db.get(JPCounter, FACTORY)
    historical_max = db.scalar(select(func.max(JPExport.last_number)).where(JPExport.factory_id == FACTORY)) or 0
    previous = max(counter.last_number if counter else 0, historical_max, preview["max_seq"])
    last = previous + len(rows)
    if last > 99999:
        raise HTTPException(409, "原编号格式的流水号已满，需先确认新的编号规则")
    now = datetime.now(ZoneInfo("Asia/Shanghai"))
    prefix = "01" + str(now.year)[2:]
    content = workbook(rows, previous + 1, prefix, creator_name or actor, preview["inputs"].get("merchandiser", ""))
    if counter:
        counter.last_number = last
    else:
        db.add(JPCounter(factory_id=FACTORY, last_number=last))
    source_ids = [m["id"] for _, m, _ in references]
    originals = [*files] + ([("粘贴正文.html", html.encode())] if html else [])
    for name, original in originals:
        checksum = sha256(original).hexdigest()
        source = db.scalar(select(Source).where(Source.factory_id == FACTORY, Source.customer_code == CUSTOMER,
            Source.kind == "jp-mail", Source.sha256 == checksum))
        if not source:
            source = Source(id=uuid4().hex, factory_id=FACTORY, customer_code=CUSTOMER,
                kind="jp-mail", sha256=checksum, file_name=Path(name).name[:240], content=original)
            db.add(source)
            db.flush()
        source_ids.append(source.id)
    record = JPExport(id=uuid4().hex, factory_id=FACTORY, request_hash=key, actor=actor,
        created_at=now.isoformat(timespec="seconds"), first_number=previous + 1, last_number=last,
        data={"rows": rows, "source_ids": source_ids, "references": current, "reason": reason, "overrides": overrides,
              "preview_fingerprint": preview["fingerprint"], "inputs": preview["inputs"], "confirmation_usd_hkd_rate": "7.8",
              "prefix": prefix, "sha256": sha256(content).hexdigest(), "warnings": preview["warnings"]}, content=content)
    db.add(record)
    db.flush()
    for identity in identities:
        db.add(JPExportIdentity(factory_id=FACTORY, identity_key=identity, export_id=record.id))
    db.flush()
    return record
