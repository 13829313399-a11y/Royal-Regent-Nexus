"""Explicit historical cutover; deliberately independent of all PO mapping engines."""
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZipFile
from zoneinfo import ZoneInfo

import openpyxl
from fastapi import HTTPException
from sqlalchemy import select

from app.models.customer_order_ledger import OrderLedgerLine as Line, OrderLedgerIdentity as Identity
from app.models.customer_order_ledger import OrderLedgerSource as Source, OrderLedgerLineSource as LineSource
from app.services import customer_order_ledger as ledger
from app.services.auth import now_text

MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = 20000
# Headers are read by name, never by customer-specific PO rules or fixed extension offsets.
HEADERS = {
    "客出单日期": "received_date", "订单类型": "order_type", "生产单号": "production_no",
    "contractno.": "contract_no", "so#/reference": "reference_no", "p/o#:": "po_no",
    "客名": "customer_name", "产品编号": "product_no", "产品中文名称": "product_name_zh",
    "产品英文名称": "product_name_en", "數量": "quantity", "数量": "quantity", "装箱": "units_per_carton",
    "箱数": "carton_count", "国家标准": "standard", "说明书": "manual", "贴纸": "label",
    "外箱贴纸": "customer_label", "箱唛": "carton_mark", "布标规格": "fabric_label", "布标格": "fabric_label",
    "包装": "packaging", "日期码": "date_code", "验货日期": "inspection_date",
    "客要求走货期": "requested_ship_date", "备注": "note", "国家": "country", "条码": "barcode",
    "印刷要求": "printing_requirement", "联系人": "contact",
}
MARKERS = {"取消单": "cancelled", "已走货订单": "shipped", "未走货订单": "unshipped"}


def text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def key(value):
    return re.sub(r"\s+", "", text(value)).replace("：", ":").casefold()


def cell_text(cell, cached):
    value = cached.value if cell.data_type == "f" else cell.value
    result = text(value)
    # Preserve numeric identifiers explicitly displayed with leading zeroes.
    if isinstance(value, (int, float)) and re.fullmatch(r"0+", cell.number_format or ""):
        result = result.zfill(len(cell.number_format))
    return result


def parse_workbook(content: bytes) -> tuple[list[dict], list[str]]:
    if len(content) > MAX_BYTES:
        raise HTTPException(400, "排期文件不能超过 25 MB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 150 * 1024 * 1024:
                raise HTTPException(400, "排期解压后过大，请拆分后迁入")
        formulas = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=False)
        values = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, "无法读取排期，请选择可打开的 .xlsx 文件") from exc
    rows, review, warnings = [], [], []
    try:
        for sheet in formulas:
            is_item = "item" in sheet.title.casefold() or "iteam" in sheet.title.casefold()
            if not is_item and sheet.title != "正单评审表":
                continue
            if sheet.max_row > MAX_ROWS or sheet.max_column > 200:
                raise HTTPException(400, f"{sheet.title} 超过迁入范围限制，请整理后重试")
            header_cells = next(sheet.iter_rows(min_row=3, max_row=3))
            headers = [text(c.value) for c in header_cells]
            fields = [HEADERS.get(key(h)) for h in headers]
            if not all(f in fields for f in ("product_no", "quantity", "reference_no", "po_no")):
                warnings.append(f"{sheet.title} 第 3 行缺少统一排期表头，未读取")
                continue
            if any(fields.count(f) > 1 for f in ("product_no", "quantity", "reference_no", "po_no", "contract_no")):
                raise HTTPException(400, f"{sheet.title} 订单关键表头重复，请先核对")
            counts = Counter(f for f in fields if f)
            section = "unshipped"
            found_markers = []
            for r, (cells, caches) in enumerate(zip(sheet.iter_rows(min_row=4), values[sheet.title].iter_rows(min_row=4)), 4):
                raw = [cell_text(c, v) for c, v in zip(cells, caches)]
                data = {f: raw[i] for i, f in enumerate(fields) if f and counts[f] == 1}
                # Only a banner without order identity changes the section. A remark mentioning cancellation does not.
                banner = MARKERS.get(key(raw[0])) if raw else None
                if banner and not any(data.get(f) for f in ("reference_no", "po_no", "product_no")):
                    section = banner
                    found_markers.append((r, section))
                    continue
                if not any(data.get(f) for f in ("reference_no", "po_no", "contract_no", "product_no")):
                    continue
                if key(data.get("product_no")) == key("产品编号"):
                    continue
                data["reference_no"] = data.get("reference_no") or data.get("po_no") or data.get("contract_no") or ""
                if not is_item:
                    if data.get("reference_no") and data.get("product_no"):
                        review.append({"data": data, "row": r, "section": section})
                    continue
                issues, blocked = [], False
                if "模拟" in data.get("note", "") or data.get("reference_no") == "SO-HY-260801":
                    issues.append("模拟/示例行，不迁入")
                    blocked = True
                if not data.get("reference_no") or not data.get("product_no"):
                    issues.append("缺少订单参考号或产品编号，可能是库存/分组行")
                    blocked = True
                if any(len(data.get(f, "")) > 255 for f in ("reference_no", "product_no")):
                    issues.append("订单编号过长，请核对")
                    blocked = True
                try:
                    data["quantity"] = ledger.decimal_text(ledger.number(data.get("quantity")))
                except HTTPException:
                    issues.append("数量缺失、非正数或精度异常，请在原表核对")
                    blocked = True
                due = data.get("requested_ship_date", "")
                if due:
                    try:
                        date.fromisoformat(due)
                    except ValueError:
                        issues.append("交期不是明确日期，请在原表核对")
                        blocked = True
                missing_cache = [headers[i] for i, c in enumerate(cells[:len(headers)])
                                 if c.data_type == "f" and caches[i].value is None and headers[i]]
                if missing_cache:
                    issues.append("公式缺少已计算值，请用 Excel 重算并保存：" + "、".join(missing_cache[:8]))
                    blocked = True
                if any(c.data_type == "e" or v.data_type == "e" for c, v in zip(cells, caches)):
                    issues.append("存在 Excel 错误值，请先核对原表")
                    blocked = True
                data["history_fields"] = [{"column": c.column_letter, "header": headers[i], "value": raw[i]}
                                          for i, c in enumerate(cells[:len(headers)]) if headers[i] and raw[i]]
                data["lineage"] = {f: f"历史排期 · {sheet.title} · 第 {r} 行" for f in data if f not in {"history_fields", "lineage"}}
                rows.append({"id": ledger.digest([sheet.title, r]), "sheet": sheet.title, "row": r,
                             "section": section, **{f: data.get(f, "") for f in ("reference_no", "product_no", "quantity", "requested_ship_date", "customer_name")},
                             "data": data, "issues": issues, "blocked": blocked, "existing": False})
            if is_item and not found_markers:
                warnings.append(f"{sheet.title} 未找到明确的取消/已走货分区，状态须逐单核对")
        if not rows:
            raise HTTPException(400, "未找到 ITEM 表订单行；请使用第 3 行含统一表头的排期文件")
        # The review sheet can be stale or contain stock/group rows. Report disagreement, never silently transfer status.
        review_index = defaultdict(list)
        for item in review:
            d = item["data"]
            review_index[(d["reference_no"].casefold(), d["product_no"].casefold())].append(item)
        for row in rows:
            matches = review_index[(row["reference_no"].casefold(), row["product_no"].casefold())]
            if len(matches) == 1:
                item = matches[0]
                row["data"]["history_review"] = {"sheet": "正单评审表", "row": item["row"], "section": item["section"], "quantity": item["data"].get("quantity", "")}
                if item["section"] != row["section"]:
                    row["issues"].append(f"正单评审表第 {item['row']} 行与 ITEM 分区不一致，请核对状态和期初走货")
                try:
                    if ledger.number(item["data"].get("quantity")) != ledger.number(row["quantity"]):
                        row["issues"].append("正单评审表与 ITEM 数量不一致，须先在原表核对")
                        row["blocked"] = True
                except HTTPException:
                    row["issues"].append("正单评审表数量无法核对，请核对原表")
            elif len(matches) > 1:
                row["issues"].append("正单评审表存在多个同号同货号行，请核对交付批次")
            else:
                row["issues"].append("正单评审表未找到唯一对应行，请核对订单范围")
        warnings.insert(0, "分区只作提示；逐单核对订单总量、状态及截至迁入日期的累计已走货量。历史期初不生成出货单，也不自动向下游发送。")
        return rows, warnings
    finally:
        formulas.close()
        values.close()


def preview(db, content, filename, factory, customer, customer_name, *, workbook=None):
    if Path(filename).suffix.lower() != ".xlsx" or Path(filename).name.startswith("._"):
        raise HTTPException(400, "请选择真实的 .xlsx 排期文件")
    rows, warnings = workbook if workbook is not None else parse_workbook(content)
    identities = defaultdict(list)
    existing = {item.identity_key: item.line_id for item in db.scalars(select(Identity).where(Identity.factory_id == factory, Identity.customer_code == customer))}
    for line in db.scalars(select(Line).where(Line.factory_id == factory, Line.customer_code == customer)):
        existing[line.identity_key] = line.id
    for row in rows:
        try:
            identity = ledger.identity_for(row["data"], customer)
            row["identity"] = identity
            identities[identity].append(row)
            if identity in existing:
                row["existing"] = True
                row["issues"].append("台账已有此订单，不覆盖；若要补齐 PO 请使用关联已有订单")
        except HTTPException as exc:
            row["blocked"] = True
            row["issues"].append(exc.detail)
    for group in identities.values():
        if len(group) > 1:
            for row in group:
                row["blocked"] = True
                row["issues"].append("文件中同订单、同货号重复，可能有多批交付；请先整理，不能直接相加或覆盖")
    return {"factory_id": factory, "customer_code": customer, "customer_name": customer_name,
            "file_name": filename, "fingerprint": ledger.digest(["history-v1", factory, customer, ledger.sha256(content).hexdigest()]),
            "rows": rows, "warnings": warnings,
            "summary": {"total": len(rows), "blocked": sum(r["blocked"] for r in rows), "existing": sum(r["existing"] for r in rows)}}


def opening_number(value):
    if str(value).strip() in {"0", "0.0", "0.00", "0.000", "0.0000"}:
        return Decimal(0)
    return ledger.number(value)


def confirm(db, *, parsed, content, selections, fingerprint, cutoff_date, reason, actor):
    if fingerprint != parsed["fingerprint"]:
        raise HTTPException(409, "文件或厂区/客户已改变，请重新预览")
    if cutoff_date > datetime.now(ZoneInfo("Asia/Shanghai")).date():
        raise HTTPException(400, "历史期初截止日期不能晚于今天")
    rows = {r["id"]: r for r in parsed["rows"]}
    if len({s.id for s in selections}) != len(selections):
        raise HTTPException(400, "迁入明细不能重复选择")
    factory, customer = parsed["factory_id"], parsed["customer_code"]
    file_hash = ledger.sha256(content).hexdigest()
    source = db.scalar(select(Source).where(Source.factory_id == factory, Source.customer_code == customer, Source.sha256 == file_hash, Source.kind == "schedule"))
    if source is None:
        source = Source(id=ledger.uid(), factory_id=factory, customer_code=customer, sha256=file_hash,
                        file_name=parsed["file_name"].replace("\\", "/").split("/")[-1][:255], kind="schedule", content=content)
        db.add(source)
        db.flush()
    created, existing_count, lines = 0, 0, []
    for selection in selections:
        row = rows.get(selection.id)
        if row is None or row["blocked"]:
            raise HTTPException(400, "所选明细不存在或仍有阻断异常，请重新核对")
        qty, opening = ledger.number(row["quantity"]), opening_number(selection.opening_shipped_quantity)
        if opening > qty:
            raise HTTPException(400, f"{row['reference_no']} 期初走货不能超过订单总量")
        record_key = ledger.digest([file_hash, row["id"], cutoff_date.isoformat(), selection.status, ledger.decimal_text(opening)])
        alias = db.get(Identity, (factory, customer, row["identity"]))
        line = db.get(Line, alias.line_id) if alias else db.scalar(select(Line).where(Line.factory_id == factory, Line.customer_code == customer, Line.identity_key == row["identity"]))
        if line:
            if line.data.get("history_migration", {}).get("record_key") != record_key:
                raise HTTPException(409, f"{row['reference_no']} 已存在，历史迁入不能覆盖已有订单")
            existing_count += 1
            lines.append(line)
            continue
        baseline = {"record_key": record_key, "source_id": source.id, "sheet": row["sheet"], "row": row["row"],
                    "section": row["section"], "cutoff_date": cutoff_date.isoformat(), "opening_shipped_quantity": ledger.decimal_text(opening),
                    "original_opening_shipped_quantity": ledger.decimal_text(opening), "confirmed_status": selection.status,
                    "reason": reason, "actor": actor, "confirmed_at": now_text()}
        data = deepcopy(row["data"])
        data.update(source_kind="schedule", history_migration=baseline)
        line = Line(id=ledger.uid(), factory_id=factory, customer_code=customer, customer_name=parsed["customer_name"],
                    identity_key=row["identity"], reference_no=row["reference_no"], product_no=row["product_no"], quantity=qty,
                    shipped_quantity=opening, status=selection.status, version=1, revision=1, data=data, created_at=now_text(), updated_at=now_text())
        db.add(line)
        db.flush()
        db.add(Identity(factory_id=factory, customer_code=customer, identity_key=row["identity"], line_id=line.id))
        db.add(LineSource(line_id=line.id, source_id=source.id))
        ledger.record_version(db, line, "历史排期迁入：" + reason, actor)
        created += 1
        lines.append(line)
    db.flush()
    return {"items": [ledger.line_out(db, line) for line in lines], "created_count": created, "existing_count": existing_count}


def correct_opening(db, line, body, actor):
    baseline = deepcopy(line.data.get("history_migration"))
    if not baseline:
        raise HTTPException(400, "此订单没有历史期初记录")
    ledger.lock_line(db, line, body.expected_revision)
    old, new = opening_number(baseline["opening_shipped_quantity"]), opening_number(body.opening_shipped_quantity)
    total = line.shipped_quantity - old + new
    if line.quantity is None or total > line.quantity or total < 0:
        raise HTTPException(400, "更正后的期初加后续走货不能超过订单总量")
    baseline.update(opening_shipped_quantity=ledger.decimal_text(new), correction_reason=body.reason, corrected_by=actor, corrected_at=now_text())
    line.data = {**line.data, "history_migration": baseline}
    line.shipped_quantity = total
    line.version += 1
    ledger.record_version(db, line, "历史期初更正：" + body.reason, actor)
    ledger.publish(db, line, ledger._existing_recipients(db, line), actor, "历史期初更正：" + body.reason)
    db.flush()
