from __future__ import annotations

import io
import math
import re
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, BinaryIO, Iterable

import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from .new_order_excel import create_new_order_workbook


FIELD_TITLES = {
    "customer_type": "客户类型", "order_date": "客出单日期", "customer_po": "客人PO",
    "huaxing_po": "华兴PO", "contract_no": "主合同号 / S/C NO", "so_no": "SO.NO",
    "customer": "客名/国家", "country": "国家", "item_no": "产品编号", "product_name": "产品名称",
    "quantity": "数量", "case_pack": "装箱", "cartons": "箱数", "shipped_quantity": "走货数量",
    "remaining_quantity": "剩余未走数量", "date_code": "生产日期码", "inspection_date": "验货期",
    "ship_date": "客要求走货期", "packaging": "包装要求", "standards": "国家标准",
    "unit_price": "单价HKD", "amount": "金额HKD", "notes": "备注", "status": "状态",
}
EXPORT_FIELDS = list(FIELD_TITLES)

ALIASES = {
    "客出单日期": "order_date", "收单日期": "order_date", "日期": "order_date",
    "客人po": "customer_po", "po号码": "customer_po", "po.no": "customer_po", "扣数po#": "customer_po",
    "华兴po": "huaxing_po", "s/cno": "contract_no", "合同": "contract_no", "主合同号": "contract_no",
    "so.no": "so_no", "so#": "so_no", "客名/国家": "customer", "客名/國家": "customer",
    "客名": "customer", "客人": "customer", "国家": "country", "產品編號": "item_no",
    "产品编号": "item_no", "产品型号": "item_no", "item#": "item_no", "產品名称": "product_name",
    "產品名稱": "product_name", "产品名称": "product_name", "产品名称/品牌": "product_name",
    "數量": "quantity", "数量": "quantity", "装箱": "case_pack", "裝箱": "case_pack",
    "箱数": "cartons", "走货数量": "shipped_quantity", "剩余未走数量": "remaining_quantity",
    "生产日期码": "date_code", "日期码": "date_code", "日期格式": "date_code",
    "验货期": "inspection_date", "客验货期": "inspection_date", "验货期(华兴回复)": "inspection_date",
    "客要求走货期": "ship_date", "走货期": "ship_date", "包装": "packaging",
    "包装要求": "packaging", "包装类别": "packaging", "国家标准": "standards",
    "单价": "unit_price", "单价hkd": "unit_price", "金额": "amount", "金额hkd": "amount",
    "备注": "notes", "生产备注": "notes", "上系统": "status", "是否已入系统": "status", "验货结果": "status",
}


def normalize_label(value: Any) -> str:
    return re.sub(r"[\s_：:]", "", str(value or "").strip().lower().replace("\n", ""))


def normalize_text(value: Any) -> str | None:
    value = re.sub(r"\s+", " ", str(value or "")).strip()
    return value or None


def coerce_number(value: Any) -> float | int | None:
    if value in (None, "", "-") or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        match = re.search(r"-?[\d,.]+", str(value))
        if not match:
            return None
        result = float(match.group().replace(",", ""))
    if not math.isfinite(result):
        return None
    return int(result) if result.is_integer() else round(result, 4)


def normalize_date(value: Any, datemode: int | None = None) -> str | None:
    if value in (None, "", "-"):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and datemode is not None and value > 20000:
        return xlrd.xldate_as_datetime(value, datemode).date().isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%m/%d/%Y", "%d/%m/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def parse_date(value: Any) -> date | None:
    try:
        return date.fromisoformat(str(value)) if value else None
    except ValueError:
        return None


def previous_workday(target: date, days: int = 7) -> date:
    result = target - timedelta(days=days)
    while result.weekday() >= 5:
        result -= timedelta(days=1)
    return result


def infer_customer_type(row: dict[str, Any], filename: str = "") -> str:
    sample = " ".join(str(row.get(key) or "") for key in ("customer_po", "huaxing_po", "contract_no", "customer"))
    return "EDU" if "EDUHX" in sample.upper() or "EDU" in filename.upper() else "彩星"


def add_derived_fields(row: dict[str, Any], today: date | None = None) -> dict[str, Any]:
    for field in ("order_date", "inspection_date", "ship_date"):
        row[field] = normalize_date(row.get(field))
    for field in ("quantity", "case_pack", "cartons", "shipped_quantity", "remaining_quantity", "unit_price", "amount"):
        row[field] = coerce_number(row.get(field))
    for field in EXPORT_FIELDS:
        if field not in {"order_date", "inspection_date", "ship_date", "quantity", "case_pack", "cartons", "shipped_quantity", "remaining_quantity", "unit_price", "amount"}:
            row[field] = normalize_text(row.get(field))
    qty, pack, shipped = row.get("quantity"), row.get("case_pack"), row.get("shipped_quantity")
    if row.get("remaining_quantity") is None and qty is not None:
        row["remaining_quantity"] = max(0, float(qty) - float(shipped or 0))
    if row.get("cartons") is None and qty is not None and pack:
        row["cartons"] = round(float(qty) / float(pack), 2)
    if row.get("amount") is None and qty is not None and row.get("unit_price") is not None:
        row["amount"] = round(float(qty) * float(row["unit_price"]), 2)
    if not row.get("inspection_date") and row.get("ship_date") and row.get("customer_type") == "EDU":
        ship = parse_date(row["ship_date"])
        if ship:
            row["inspection_date"] = previous_workday(ship).isoformat()
    if not row.get("status"):
        row["status"] = "已走货" if row.get("remaining_quantity") == 0 else "未走货"
    row["flags"] = build_flags(row, today)
    row["risk_level"] = "high" if any(x["level"] == "high" for x in row["flags"]) else "medium" if row["flags"] else "normal"
    return row


def build_flags(row: dict[str, Any], today: date | None = None) -> list[dict[str, str]]:
    result: list[dict[str, str]] = []
    for field, message in {"customer_po": "缺少客人PO", "item_no": "缺少产品编号", "product_name": "缺少产品名称", "quantity": "缺少数量", "ship_date": "缺少走货期"}.items():
        if row.get(field) in (None, "", 0):
            result.append({"level": "high", "code": f"missing_{field}", "text": message})
    if not row.get("case_pack"):
        result.append({"level": "medium", "code": "missing_case_pack", "text": "装箱数待确认"})
    if not row.get("date_code"):
        result.append({"level": "medium", "code": "missing_date_code", "text": "日期码待确认"})
    qty, pack = row.get("quantity"), row.get("case_pack")
    if qty and pack and float(qty) % float(pack):
        result.append({"level": "medium", "code": "partial_carton", "text": "数量不能整除装箱数"})
    ship, remaining, ref = parse_date(row.get("ship_date")), float(row.get("remaining_quantity") or 0), today or date.today()
    if ship and remaining > 0 and ship < ref:
        result.append({"level": "high", "code": "overdue", "text": "未走货订单已逾期"})
    elif ship and remaining > 0 and 0 <= (ship - ref).days <= 14:
        result.append({"level": "medium", "code": "ship_soon", "text": "14天内走货"})
    return result


def _header(rows: list[list[Any]]) -> tuple[int, dict[int, str]]:
    best = (-1, {})
    for index, row in enumerate(rows[:15]):
        mapping: dict[int, str] = {}
        for column, value in enumerate(row):
            field = ALIASES.get(normalize_label(value))
            if field and field not in mapping.values():
                mapping[column] = field
        if len(mapping) > len(best[1]):
            best = (index, mapping)
    return best


def _parse_rows(rows: list[list[Any]], sheet: str, filename: str, datemode: int | None) -> list[dict[str, Any]]:
    header_index, mapping = _header(rows)
    if header_index < 0 or len(mapping) < 4:
        return []
    result = []
    for row_no, values in enumerate(rows[header_index + 1:], header_index + 2):
        row = {field: values[col] if col < len(values) else None for col, field in mapping.items()}
        for field in ("order_date", "inspection_date", "ship_date"):
            row[field] = normalize_date(row.get(field), datemode)
        if not row.get("item_no") and not row.get("product_name"):
            continue
        if not any(row.get(key) not in (None, "") for key in ("customer_po", "huaxing_po", "contract_no", "customer")):
            continue
        if str(row.get("item_no") or "").strip() in {"合计", "总计"}:
            continue
        row.update(customer_type=infer_customer_type(row, filename), source_sheet=sheet, source_row=row_no)
        result.append(add_derived_fields(row))
    return result


def read_schedule(source: str | Path | bytes | BinaryIO, filename: str | None = None, **_: Any) -> dict[str, Any]:
    data = source if isinstance(source, bytes) else source.read() if hasattr(source, "read") else Path(source).read_bytes()
    name = filename or (Path(source).name if isinstance(source, (str, Path)) else "schedule.xlsx")
    sheets: list[tuple[str, list[list[Any]], int | None]] = []
    if bytes(data[:8]) == bytes.fromhex("D0CF11E0A1B11AE1"):
        book = xlrd.open_workbook(file_contents=data)
        for ws in book.sheets():
            sheets.append((ws.name, [[ws.cell_value(r, c) for c in range(ws.ncols)] for r in range(ws.nrows)], book.datemode))
    else:
        book = load_workbook(io.BytesIO(data), data_only=True)
        for ws in book.worksheets:
            sheets.append((ws.title, [[cell.value for cell in row] for row in ws.iter_rows()], None))
    priority = {"ITEM表": 0, "Iteam表": 0, "正单评审表": 1, "接单表": 2, "PO订单": 3, "订单": 3}
    sheets.sort(key=lambda item: (priority.get(item[0], 9), item[0]))
    best: tuple[str, list[dict[str, Any]]] | None = None
    for sheet, rows, datemode in sheets:
        records = _parse_rows(rows, sheet, name, datemode)
        if records and (best is None or len(records) > len(best[1])):
            best = (sheet, records)
        if records and priority.get(sheet) == 0:
            best = (sheet, records)
            break
    if not best:
        raise ValueError("Excel 中未找到可识别的华兴排期表头。")
    sheet, records = best
    return {"filename": name, "sheet": sheet, "records": records, "summary": summarize_records(records),
            "meta": {"sheet_names": [item[0] for item in sheets], "rows": len(records)}}


def summarize_records(records: Iterable[dict[str, Any]]) -> dict[str, Any]:
    rows = list(records)
    customers = Counter(str(row.get("customer") or "待确认") for row in rows)
    return {"orders": len(rows), "po_count": len({row.get("customer_po") for row in rows if row.get("customer_po")}),
            "quantity": round(sum(float(row.get("quantity") or 0) for row in rows), 2),
            "remaining": round(sum(float(row.get("remaining_quantity") or 0) for row in rows), 2),
            "amount": round(sum(float(row.get("amount") or 0) for row in rows), 2),
            "high_risk": sum(row.get("risk_level") == "high" for row in rows),
            "flagged": sum(row.get("risk_level") != "normal" for row in rows), "customers": customers.most_common(8)}


def _style(ws) -> None:
    fill = PatternFill("solid", fgColor="17324D")
    for cell in ws[1]:
        cell.fill, cell.font = fill, Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
    for column in range(1, ws.max_column + 1):
        width = max(len(str(ws.cell(row, column).value or "")) for row in range(1, min(ws.max_row, 250) + 1))
        ws.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 10), 32)


def create_import_workbook(
    records: list[dict[str, Any]],
    output_path: Path,
    existing: list[dict[str, Any]] | None = None,
    template_source: str | Path | bytes | BinaryIO | None = None,
) -> Path:
    if template_source is not None:
        prepared = [add_derived_fields(dict(source)) for source in records]
        aliases = {
            field: tuple(
                dict.fromkeys(
                    [title] + [alias for alias, target in ALIASES.items() if target == field]
                )
            )
            for field, title in FIELD_TITLES.items()
        }
        create_new_order_workbook(
            template_source,
            output_path,
            prepared,
            aliases,
            filename=Path(template_source).name if isinstance(template_source, (str, Path)) else "schedule.xlsx",
            sheet_title="新单",
        )
        return output_path
    wb = Workbook(); ws = wb.active; ws.title = "AI导入待确认"
    ws.append([FIELD_TITLES[x] for x in EXPORT_FIELDS] + ["风险等级", "校验结果"])
    for source in records:
        row = add_derived_fields(dict(source))
        ws.append([row.get(x) for x in EXPORT_FIELDS] + [row.get("risk_level"), "；".join(x["text"] for x in row.get("flags", []))])
    _style(ws)
    check = wb.create_sheet("校验汇总"); check.append(["级别", "PO", "产品编号", "问题"])
    for row in records:
        for flag in row.get("flags", []):
            check.append([flag["level"], row.get("customer_po"), row.get("item_no"), flag["text"]])
    _style(check)
    if existing:
        history = wb.create_sheet("原排期数据"); history.append([FIELD_TITLES[x] for x in EXPORT_FIELDS])
        for row in existing:
            history.append([row.get(x) for x in EXPORT_FIELDS])
        _style(history)
    output_path.parent.mkdir(parents=True, exist_ok=True); wb.save(output_path)
    return output_path


def create_summary_workbook(records: list[dict[str, Any]], output_path: Path) -> Path:
    return create_import_workbook(records, output_path)
