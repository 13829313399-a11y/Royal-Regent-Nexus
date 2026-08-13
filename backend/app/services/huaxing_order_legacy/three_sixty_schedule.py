from __future__ import annotations

import io
import math
import re
import zipfile
from copy import copy
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from functools import lru_cache
from pathlib import Path
from typing import Any, BinaryIO, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.cell_range import CellRange

from .new_order_excel import (
    _copy_style,
    _extend_subtotal_to_previous_row,
    _rewrite_formula_for_insert,
    _shift_target_sheet_structures,
    load_complete_workbook_compatible,
    load_workbook_compatible,
)


FIELD_TITLES: dict[str, str] = {
    "system_status": "系统状态", "order_date": "出单日期", "po_no": "PO号",
    "customer_release_no": "客户生产单号", "production_no": "360生产单号",
    "customer_po": "客户PO号", "contract_no": "360合同号", "customer": "客名",
    "item_no": "款号", "item_full": "产品货号", "ship_item_no": "出货货号",
    "product_name": "产品名称", "version": "版本", "spec": "规格",
    "quantity": "PO数量", "outer_pack": "外箱装箱数", "cartons": "箱数",
    "manual": "说明书", "artwork": "彩盒", "vehicle": "车款",
    "date_code": "日期码", "special_remark": "备注", "inspection_date": "验货日期",
    "customer_label": "PO利宝", "inspection_result": "验货结果",
    "third_party_inspection": "第三方公证行验货",
    "fcd_date": "FCD期", "po_ship_date": "走货期", "merchandiser": "跟单",
    "unit_price_usd": "订单单价USD", "unit_price_hkd": "订单单价HK$",
    "total_usd": "总金额USD", "total_hkd": "总金额HK$",
    "customer_country": "走货国家", "barcode": "条码", "factory_no": "工厂编号",
    "container_type": "MS Container Type (FCL/LCL)",
    "consolidated_id": "Consolidated ID", "port_of_discharge": "Port of Discharge",
    "transportation_mode": "走货方式", "rl_order_type": "系统",
    "release_revision": "Release Revision",
}
EXPORT_FIELDS = list(FIELD_TITLES)
DATE_FIELDS = {"order_date", "inspection_date", "fcd_date", "po_ship_date"}
NUMBER_FIELDS = {"quantity", "outer_pack", "cartons", "unit_price_usd", "unit_price_hkd", "total_usd", "total_hkd"}

FIELD_ALIASES = {
    "系统状态": "system_status", "出单日期": "order_date", "来单日期": "order_date",
    "入单日期": "order_date",
    "PO号": "po_no", "美国PO号": "po_no", "生产单号": "customer_release_no",
    "美国生产单号": "customer_release_no", "360生产单号": "production_no",
    "WMPO号": "customer_po", "客户PO号": "customer_po", "360合同号": "contract_no",
    "第三方客户PONO#": "customer_po", "主合同号": "contract_no",
    # 360 当前排期的“正单合同号”实际保存 RL-xxxx-x 生产单号，而非主合同数字。
    "正单合同号": "production_no",
    "客名": "customer", "客户名称": "customer", "產品編號": "item_no",
    "款号": "item_no",
    "产品编号": "item_no", "產品型號": "item_no", "产品型号": "item_no",
    "产品货号": "item_full", "出货货号": "ship_item_no",
    "产品名称": "product_name", "產品名稱": "product_name", "版本": "version", "规格": "spec",
    "產品規格": "spec", "PO数量": "quantity", "数量": "quantity",
    "外箱装箱数": "outer_pack", "裝箱隻數": "outer_pack", "装箱": "outer_pack",
    "箱数": "cartons", "总箱数": "cartons", "说明书": "manual", "彩盒": "artwork",
    "车款": "vehicle", "日期码": "date_code", "备注": "special_remark",
    "PO利宝": "customer_label", "外箱PO利宝": "customer_label", "验货结果": "inspection_result",
    "第三方公证行验货": "third_party_inspection",
    "验货日期": "inspection_date", "验货期": "inspection_date", "FCD期": "fcd_date",
    "走货期FCD": "fcd_date",
    "走货期": "po_ship_date", "PO走货期": "po_ship_date", "跟单": "merchandiser",
    "订单单价USD": "unit_price_usd", "订单单价HK$": "unit_price_hkd", "单价HK$": "unit_price_hkd",
    "单价HKD": "unit_price_hkd", "单价": "unit_price_hkd",
    "总金额USD": "total_usd", "总金额HK$": "total_hkd", "金额HKD": "total_hkd",
    "走货国家": "customer_country", "条码": "barcode", "工厂编号": "factory_no",
    "MSContainerTypeFCLLCL": "container_type", "ConsolidatedID": "consolidated_id",
    "PortofDischarge": "port_of_discharge", "走货方式": "transportation_mode",
    "系统": "rl_order_type", "ReleaseRevision": "release_revision",
}


@dataclass(frozen=True)
class SourceRef:
    kind: str
    path: Path
    member: str | None = None


def open_source_bytes(ref: SourceRef) -> bytes:
    if ref.kind == "zip":
        if not ref.member:
            raise ValueError("压缩包来源缺少文件名。")
        with zipfile.ZipFile(ref.path) as archive:
            return archive.read(ref.member)
    return ref.path.read_bytes()


def normalize_label(value: Any) -> str:
    text = str(value or "").strip().replace("（", "(").replace("）", ")")
    return re.sub(r"[\s/\\()]+", "", text)


def normalize_text(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def coerce_number(value: Any) -> float | int | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        number = float(value)
    else:
        try:
            number = float(str(value).replace(",", "").strip())
        except ValueError:
            return None
    return int(number) if number.is_integer() else number


def normalize_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and 25000 <= float(value) <= 70000:
        try:
            return from_excel(value).date().isoformat()
        except Exception:
            return None
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%b-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text or None


def normalize_value(field: str, value: Any) -> Any:
    if field in DATE_FIELDS:
        return normalize_date(value)
    if field in NUMBER_FIELDS:
        return coerce_number(value)
    if isinstance(value, (date, datetime)):
        return normalize_date(value)
    return normalize_text(value)


def parse_iso_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10]) if value else None
    except ValueError:
        return None


def build_360_date_code(value: Any, factory_no: Any = "60350") -> str | None:
    target = parse_iso_date(value)
    if not target:
        return None
    factory = re.sub(r"\D", "", str(factory_no or ""))[:5] or "60350"
    # 跟单提供的“日期码示例”明确：2026-07-02 -> 60350A26183，
    # 即工厂号 + A + 两位年份 + 当年第几日，不再错误倒推 15 个工作日。
    return f"{factory}A{target.year % 100:02d}{target.timetuple().tm_yday:03d}"


def contains_pending(value: Any) -> bool:
    text = str(value or "")
    return any(token in text for token in ("欠", "待", "未", "TBD", "订购"))


def build_flags(record: dict[str, Any], today: date | None = None) -> list[dict[str, str]]:
    today = today or date.today()
    flags: list[dict[str, str]] = []
    required = [("production_no", "360生产单号"), ("contract_no", "合同号"),
                ("item_no", "产品编号"), ("quantity", "PO数量"),
                ("inspection_date", "验货日期")]
    if record.get("source_sheet") != "接单表":
        required.extend((("customer_release_no", "客户Release号"), ("merchandiser", "跟单"),
                         ("unit_price_usd", "订单单价USD")))
    for field, label in required:
        if not record.get(field):
            flags.append({
                "level": "medium", "code": f"missing_{field}",
                "field": field, "text": f"缺{label}",
            })
    expected = record.get("expected_date_code") or build_360_date_code(
        record.get("inspection_date") or record.get("fcd_date"),
        record.get("factory_no"),
    )
    actual = str(record.get("date_code") or "").strip()
    if expected and not actual:
        flags.append({
            "level": "medium", "code": "missing_date_code", "field": "date_code",
            "text": f"缺日期码，建议{expected}",
        })
    elif expected and actual != expected:
        flags.append({
            "level": "medium", "code": "date_code_mismatch", "field": "date_code",
            "text": f"日期码应复核，建议{expected}",
        })
    qty, outer, cartons = (coerce_number(record.get(key)) for key in ("quantity", "outer_pack", "cartons"))
    if qty is not None and not outer:
        flags.append({
            "level": "medium", "code": "missing_outer_pack", "field": "outer_pack",
            "text": "缺外箱装箱数",
        })
    if qty is not None and outer and cartons is not None and abs(cartons - qty / outer) > .02:
        flags.append({
            "level": "high", "code": "carton_count_mismatch", "field": "cartons",
            "text": "箱数与PO数量/装箱数不一致",
        })
    inspection = parse_iso_date(record.get("inspection_date"))
    ship = parse_iso_date(record.get("po_ship_date"))
    if inspection and inspection < today and not record.get("inspection_result"):
        flags.append({
            "level": "high", "code": "missing_inspection_result",
            "field": "inspection_result", "text": "验货日期已过，结果未登记",
        })
    if ship and inspection and ship < inspection:
        flags.append({
            "level": "high", "code": "ship_before_inspection", "field": "po_ship_date",
            "text": "走货期早于验货日期",
        })
    if not ship and record.get("source_sheet") != "接单表":
        flags.append({
            "level": "medium", "code": "missing_po_ship_date", "field": "po_ship_date",
            "text": "缺走货期",
        })
    if contains_pending(record.get("special_remark")):
        flags.append({"level": "medium", "text": "备注含待办事项"})
    return flags


def add_derived_fields(record: dict[str, Any]) -> dict[str, Any]:
    qty, outer, cartons = (coerce_number(record.get(key)) for key in ("quantity", "outer_pack", "cartons"))
    if cartons is None and qty is not None and outer:
        calculated = qty / outer
        record["cartons"] = int(calculated) if calculated.is_integer() else round(calculated, 2)
    usd, hkd = coerce_number(record.get("unit_price_usd")), coerce_number(record.get("unit_price_hkd"))
    if hkd is None and usd is not None:
        hkd = round(usd * 7.75, 4)
        record["unit_price_hkd"] = hkd
    if qty is not None and usd is not None and coerce_number(record.get("total_usd")) is None:
        record["total_usd"] = round(qty * usd, 2)
    if qty is not None and hkd is not None and coerce_number(record.get("total_hkd")) is None:
        record["total_hkd"] = round(qty * hkd, 2)
    calculated_code = build_360_date_code(
        record.get("inspection_date") or record.get("fcd_date"),
        record.get("factory_no"),
    )
    record["calculated_date_code"] = calculated_code
    if record.get("date_code_basis") == "接单表同验货期历史日期码" and record.get("date_code"):
        record["expected_date_code"] = record["date_code"]
    else:
        record["expected_date_code"] = calculated_code
    record["flags"] = build_flags(record)
    record["risk_level"] = "high" if any(flag["level"] == "high" for flag in record["flags"]) else ("medium" if record["flags"] else "normal")
    return record


def find_header_row(ws) -> tuple[int, dict[int, str]]:
    best: tuple[int, dict[int, str]] | None = None
    for row_idx in range(1, min(ws.max_row, 15) + 1):
        mapping: dict[int, str] = {}
        for col_idx in range(1, min(ws.max_column, 100) + 1):
            field = FIELD_ALIASES.get(normalize_label(ws.cell(row_idx, col_idx).value))
            if field:
                mapping[col_idx] = field
        if best is None or len(mapping) > len(best[1]):
            best = (row_idx, mapping)
    if not best or len(best[1]) < 6 or not {"item_no", "product_name", "quantity"}.issubset(best[1].values()):
        raise ValueError(f"工作表“{ws.title}”未找到可识别的360排期表头。")
    return best


def read_schedule(source: str | Path | bytes | BinaryIO, *, filename: str | None = None,
                  sheet_name: str | None = None, **_: Any) -> dict[str, Any]:
    raw = io.BytesIO(source) if isinstance(source, bytes) else source
    workbook = load_workbook(raw, data_only=False)
    if isinstance(source, bytes):
        value_workbook = load_workbook(io.BytesIO(source), data_only=True)
    else:
        value_workbook = load_workbook(source, data_only=True)
    if sheet_name and sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
    else:
        candidates = []
        for candidate in workbook.worksheets:
            try:
                header, mapping = find_header_row(candidate)
                candidates.append((len(mapping) + (20 if "Iteam" in candidate.title else 0), header, candidate))
            except ValueError:
                pass
        if not candidates:
            raise ValueError("Excel中未找到可识别的360排期工作表。")
        ws = max(candidates, key=lambda item: item[0])[2]
    header_row, mapping = find_header_row(ws)
    records: list[dict[str, Any]] = []
    reference_pattern = re.compile(r"^='?([^']+)'?!([A-Z]+\d+)$|^=([^!]+)!([A-Z]+\d+)$")
    def resolved_cell(row_idx: int, col_idx: int) -> Any:
        value = ws.cell(row_idx, col_idx).value
        if not isinstance(value, str) or not value.startswith("="):
            return value
        match = reference_pattern.match(value)
        if match:
            target_sheet = match.group(1) or match.group(3)
            target_cell = match.group(2) or match.group(4)
            if target_sheet in value_workbook.sheetnames:
                return value_workbook[target_sheet][target_cell].value
        # 普通公式（如 =G5*H5）回退到缓存计算值，不能一律当 None 丢数据。
        if ws.title in value_workbook.sheetnames:
            return value_workbook[ws.title].cell(row_idx, col_idx).value
        return None
    for row_idx in range(header_row + 1, ws.max_row + 1):
        record = {field: normalize_value(field, resolved_cell(row_idx, col_idx)) for col_idx, field in mapping.items()}
        if not any(record.get(key) for key in ("production_no", "contract_no", "customer", "item_no", "product_name")):
            continue
        identity_values = [str(record.get(key) or "").strip() for key in ("production_no", "contract_no", "customer", "item_no", "product_name")]
        if any(value in {"合计", "总计", "小计"} or value.startswith(("合计：", "总计：", "小计：")) for value in identity_values):
            continue
        record.update({"row_number": row_idx, "source_sheet": ws.title})
        records.append(record)
    historical_codes: dict[str, str] = {}
    grouped_codes: dict[str, Counter[str]] = {}
    for record in records:
        inspection = record.get("inspection_date") or record.get("fcd_date")
        code = str(record.get("date_code") or "").strip()
        if inspection and code:
            grouped_codes.setdefault(str(inspection), Counter())[code] += 1
    for inspection, counts in grouped_codes.items():
        historical_codes[inspection] = counts.most_common(1)[0][0]
    for record in records:
        inspection = str(record.get("inspection_date") or record.get("fcd_date") or "")
        if inspection in historical_codes:
            record["date_code_basis"] = "接单表同验货期历史日期码"
            record["expected_date_code"] = historical_codes[inspection]
        add_derived_fields(record)
    result = {"filename": filename or "schedule.xlsx", "sheet": ws.title,
              "meta": {"sheet": ws.title, "header_row": header_row, "columns": [FIELD_TITLES[field] for field in mapping.values()]},
              "records": records, "summary": summarize_records(records), "warnings": []}
    workbook.close()
    value_workbook.close()
    return result


def _item_digits(value: Any) -> str:
    return re.sub(r"\D", "", str(value or ""))


def _full_item_key(value: Any) -> str:
    text = re.sub(r"\s+", "", str(value or "")).upper()
    return re.sub(r"\.0+$", "", text)


@lru_cache(maxsize=4)
def _schedule_index(template: bytes) -> dict[str, Any]:
    """Build reusable schedule master data once per uploaded template.

    A 360 workbook is several megabytes.  The old request path opened and scanned
    the same workbook four times (reference data, date code, duplicate check and
    export), which made a two-file upload appear stuck.  This immutable cache is
    keyed by the exact template bytes and is naturally invalidated after a new
    schedule is uploaded.
    """
    workbook = load_workbook(io.BytesIO(template), data_only=True, read_only=False)
    existing_numbers: set[str] = set()
    full_item_to_base: dict[str, Counter[str]] = {}
    full_item_to_ship_item: dict[str, Counter[str]] = {}
    names: dict[str, Counter[str]] = {}
    countries: dict[str, Counter[str]] = {}
    prices: dict[str, Counter[float]] = {}
    historical_codes: dict[tuple[str, str], Counter[str]] = {}
    try:
        for ws in workbook.worksheets:
            try:
                header_row, mapping = find_header_row(ws)
            except ValueError:
                continue
            columns = {field: col for col, field in mapping.items()}
            for row_no in range(header_row + 1, ws.max_row + 1):
                if "production_no" in columns:
                    production = str(ws.cell(row_no, columns["production_no"]).value or "").upper()
                    match = re.search(r"RL-\d+-\d+", production)
                    if match:
                        existing_numbers.add(match.group(0))

                item = (
                    _item_digits(ws.cell(row_no, columns["item_no"]).value)
                    if "item_no" in columns
                    else ""
                )
                if item:
                    if "item_full" in columns:
                        full_key = _full_item_key(ws.cell(row_no, columns["item_full"]).value)
                        if full_key:
                            full_item_to_base.setdefault(full_key, Counter())[item] += 1
                            if "ship_item_no" in columns:
                                ship_item = _item_digits(
                                    ws.cell(row_no, columns["ship_item_no"]).value
                                )
                                if ship_item:
                                    full_item_to_ship_item.setdefault(
                                        full_key, Counter()
                                    )[ship_item] += 1
                    if "product_name" in columns:
                        value = normalize_text(ws.cell(row_no, columns["product_name"]).value)
                        if value:
                            names.setdefault(item, Counter())[value] += 1
                    if "unit_price_usd" in columns:
                        value = coerce_number(ws.cell(row_no, columns["unit_price_usd"]).value)
                        if value is not None:
                            prices.setdefault(item, Counter())[float(value)] += 1

                if "customer" in columns and "customer_country" in columns:
                    customer = normalize_text(ws.cell(row_no, columns["customer"]).value)
                    country = normalize_text(ws.cell(row_no, columns["customer_country"]).value)
                    if customer and country:
                        countries.setdefault(
                            normalize_label(customer).casefold(), Counter()
                        )[country] += 1

                if "date_code" in columns:
                    inspection = ""
                    for date_field in ("inspection_date", "fcd_date"):
                        if date_field in columns:
                            inspection = str(
                                normalize_date(ws.cell(row_no, columns[date_field]).value) or ""
                            )
                            if inspection:
                                break
                    code = str(ws.cell(row_no, columns["date_code"]).value or "").strip()
                    match = re.match(r"(\d{5})A\d{5}$", code, re.I)
                    if inspection and match:
                        historical_codes.setdefault(
                            (match.group(1), inspection), Counter()
                        )[code] += 1
    finally:
        workbook.close()
    return {
        "existing_numbers": frozenset(existing_numbers),
        "full_item_to_base": full_item_to_base,
        "full_item_to_ship_item": full_item_to_ship_item,
        "names": names,
        "countries": countries,
        "prices": prices,
        "historical_codes": historical_codes,
    }


def apply_schedule_date_codes(template: bytes, records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Prefer an established code only for the same factory and inspection/FCD date."""
    historical = _schedule_index(template)["historical_codes"]
    result: list[dict[str, Any]] = []
    for source in records:
        row = dict(source)
        inspection = str(row.get("inspection_date") or row.get("fcd_date") or "")
        factory = re.sub(r"\D", "", str(row.get("factory_no") or ""))[:5] or "60350"
        key = (factory, inspection)
        if key in historical:
            row["date_code"] = historical[key].most_common(1)[0][0]
            row["date_code_basis"] = "当前排期同工厂同验货期历史日期码"
        else:
            row["date_code"] = row.get("date_code") or build_360_date_code(inspection, factory)
            row["date_code_basis"] = "工厂号+A+年份+验货/FCD日年内序号"
        result.append(add_derived_fields(row))
    return result


def filter_existing_schedule_records(
    template: bytes,
    records: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Exclude RL rows already present in current/completed/cancelled schedule sheets."""
    existing_numbers = _schedule_index(template)["existing_numbers"]

    new_rows: list[dict[str, Any]] = []
    existing_rows: list[dict[str, Any]] = []
    for source in records:
        row = dict(source)
        match = re.search(r"RL-\d+-\d+", str(row.get("production_no") or "").upper())
        if match and match.group(0) in existing_numbers:
            existing_rows.append(row)
        else:
            new_rows.append(row)
    return new_rows, existing_rows


def apply_schedule_reference_data(
    template: bytes,
    records: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[str]]:
    """Fill only stable, unique product/customer master data from the latest schedule."""
    index = _schedule_index(template)
    full_item_to_base = index["full_item_to_base"]
    full_item_to_ship_item = index["full_item_to_ship_item"]
    names = index["names"]
    countries = index["countries"]
    prices = index["prices"]

    warnings: list[str] = []
    prepared: list[dict[str, Any]] = []
    missing_name_rows: dict[str, list[str]] = {}
    conflicting_price_rows: dict[str, list[str]] = {}
    for source in records:
        row = dict(source)
        full_key = _full_item_key(row.get("item_full"))
        if full_key and full_key in full_item_to_base:
            # 360 货号长度并不统一。旧款常取前 7 位，新款可能保留小数点前
            # 的 10 位；以当前有效排期中的同一完整货号为最高优先级。
            row["item_no"] = full_item_to_base[full_key].most_common(1)[0][0]
            if full_key in full_item_to_ship_item:
                row["ship_item_no"] = full_item_to_ship_item[full_key].most_common(1)[0][0]
            else:
                row["ship_item_no"] = row["item_no"]
            row["item_no_source"] = "最新360排期按完整货号继承"
        item = _item_digits(row.get("item_no"))
        if item and item in names:
            row["product_name"] = names[item].most_common(1)[0][0]
            row["product_name_source"] = "最新360排期按款号继承"
        customer_key = normalize_label(row.get("customer")).casefold()
        if customer_key and customer_key in countries and not row.get("customer_country"):
            row["customer_country"] = countries[customer_key].most_common(1)[0][0]
        if item and item in prices and row.get("unit_price_usd") in (None, ""):
            distinct = list(prices[item])
            if len(distinct) == 1:
                row["unit_price_usd"] = distinct[0]
                row["price_source"] = "最新360排期同款号唯一单价"
            else:
                conflicting_price_rows.setdefault(item, []).append(
                    str(row.get("production_no") or item)
                )
        if item and item not in names:
            missing_name_rows.setdefault(item, []).append(
                str(row.get("production_no") or item)
            )
        prepared.append(add_derived_fields(row))
    for item, numbers in conflicting_price_rows.items():
        warnings.append(
            f"款号 {item} 在最新排期存在多个历史单价，涉及 {len(numbers)} 份 PO，"
            "系统未猜价，请以主合同核对。"
        )
    for item, numbers in missing_name_rows.items():
        warnings.append(
            f"款号 {item} 在最新排期没有中文货名，涉及 {len(numbers)} 份 PO，"
            "已保留 PO 英文 Description；取得标准中文货名后可由排期自动继承。"
        )
    return prepared, list(dict.fromkeys(warnings))


def summarize_records(records: Iterable[dict[str, Any]]) -> dict[str, Any]:
    rows = list(records)
    return {
        "orders": len(rows), "contracts": len({str(row.get("contract_no")) for row in rows if row.get("contract_no")}),
        "quantity": round(sum(float(coerce_number(row.get("quantity")) or 0) for row in rows), 2),
        "cartons": round(sum(float(coerce_number(row.get("cartons")) or 0) for row in rows), 2),
        "high_risk": sum(row.get("risk_level") == "high" for row in rows),
        "flagged": sum(row.get("risk_level") != "normal" for row in rows),
        "customers": Counter(str(row.get("customer") or "待确认") for row in rows).most_common(8),
    }


def _style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="17324D")
    thin = Side(style="thin", color="D9E2EC")
    for cell in ws[1]:
        cell.fill, cell.font = header_fill, Font(color="FFFFFF", bold=True)
        cell.alignment, cell.border = Alignment(horizontal="center", vertical="center", wrap_text=True), Border(bottom=thin)
    ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
    for col in range(1, ws.max_column + 1):
        width = max(len(str(ws.cell(row, col).value or "")) for row in range(1, min(ws.max_row, 250) + 1))
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 10), 34)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def create_import_workbook(records: list[dict[str, Any]], output_path: Path, _template: bytes | None = None) -> Path:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "AI导入待确认"
    ws.append([FIELD_TITLES[field] for field in EXPORT_FIELDS] + ["风险等级", "检查结果", "建议日期码"])
    for source in records:
        row = add_derived_fields(dict(source))
        ws.append([row.get(field) for field in EXPORT_FIELDS] + [row["risk_level"], "；".join(flag["text"] for flag in row["flags"]), row.get("expected_date_code")])
    _style_sheet(ws)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def warm_schedule_template(template: bytes) -> None:
    """Preload reusable schedule indexes after service start/upload."""
    _schedule_index(template)


_TOTAL_LABELS = {normalize_label(value) for value in ("合计", "总计", "小计")}
_SIMPLE_SHEET_REFERENCE = re.compile(
    r"^=(?:'(?P<quoted>(?:[^']|'')+)'|(?P<plain>[^'!]+))!"
    r"(?P<column>\$?[A-Z]{1,3})(?P<row>\$?\d+)$",
    re.I,
)
_BUSINESS_FIELDS = ("production_no", "contract_no", "item_no", "customer", "quantity")
_FORMULA_FIELDS = {"cartons", "unit_price_hkd", "total_usd", "total_hkd"}
_IDENTIFIER_FIELDS = {
    "po_no", "customer_release_no", "production_no", "customer_po", "contract_no",
    "item_no", "item_full", "ship_item_no", "date_code", "barcode",
}


def _sheet_by_names(workbook, names: Iterable[str]):
    wanted = {str(name).strip().casefold() for name in names}
    return next(
        (sheet for sheet in workbook.worksheets if sheet.title.strip().casefold() in wanted),
        None,
    )


def _logical_max_column(ws, header_row: int) -> int:
    populated = [
        column
        for column in range(1, min(ws.max_column or 1, 512) + 1)
        if ws.cell(header_row, column).value not in (None, "")
    ]
    return max(populated, default=1)


def _header_columns(ws, header_row: int) -> dict[str, list[int]]:
    columns: dict[str, list[int]] = defaultdict(list)
    for column in range(1, _logical_max_column(ws, header_row) + 1):
        field = FIELD_ALIASES.get(normalize_label(ws.cell(header_row, column).value))
        if field:
            columns[field].append(column)
    return dict(columns)


def _effective_value(
    workbook,
    data_workbook,
    sheet_name: str,
    row: int,
    column: int,
    *,
    depth: int = 0,
) -> Any:
    formula_sheet = _sheet_by_names(workbook, (sheet_name,))
    data_sheet = _sheet_by_names(data_workbook, (sheet_name,))
    if formula_sheet is None:
        return None
    if data_sheet is not None:
        cached = data_sheet.cell(row, column).value
        if cached not in (None, ""):
            return cached
    raw = formula_sheet.cell(row, column).value
    if not isinstance(raw, str) or not raw.startswith("=") or depth >= 3:
        return raw
    match = _SIMPLE_SHEET_REFERENCE.fullmatch(raw)
    if not match:
        return None
    referenced_sheet = (match.group("quoted") or match.group("plain") or "").replace("''", "'").strip()
    return _effective_value(
        workbook,
        data_workbook,
        referenced_sheet,
        int(match.group("row").replace("$", "")),
        column_index_from_string(match.group("column").replace("$", "")),
        depth=depth + 1,
    )


def _find_total_row(workbook, data_workbook, ws, header_row: int) -> int:
    data_sheet = _sheet_by_names(data_workbook, (ws.title,))
    max_column = _logical_max_column(ws, header_row)
    for row in range(header_row + 1, (ws.max_row or header_row) + 1):
        for column in range(1, max_column + 1):
            values = [ws.cell(row, column).value]
            if data_sheet is not None:
                values.append(data_sheet.cell(row, column).value)
            values.append(_effective_value(workbook, data_workbook, ws.title, row, column))
            if any(
                value not in (None, "")
                and not (isinstance(value, str) and value.startswith("="))
                and normalize_label(value) in _TOTAL_LABELS
                for value in values
            ):
                return row
    raise ValueError(f"360 排期工作表“{ws.title.strip()}”未找到合计行。")


def _row_has_business_content(
    workbook,
    data_workbook,
    ws,
    row: int,
    columns: dict[str, list[int]],
) -> bool:
    for field in _BUSINESS_FIELDS:
        for column in columns.get(field, ()):
            value = _effective_value(workbook, data_workbook, ws.title, row, column)
            if value not in (None, "", "-") and normalize_label(value) not in _TOTAL_LABELS:
                return True
    return False


def _last_business_row(
    workbook,
    data_workbook,
    ws,
    header_row: int,
    total_row: int,
    columns: dict[str, list[int]],
) -> int:
    for row in range(total_row - 1, header_row, -1):
        if _row_has_business_content(workbook, data_workbook, ws, row, columns):
            return row
    return header_row


def _insert_rows_preserving_workbook(workbook, target, insert_row: int, amount: int) -> None:
    if amount <= 0:
        return
    formulas = [
        (sheet, cell.row, cell.column, cell.value)
        for sheet in workbook.worksheets
        for cell in sheet._cells.values()
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    moved_merges = [
        CellRange(str(value))
        for value in target.merged_cells.ranges
        if value.max_row >= insert_row
    ]
    for merged in moved_merges:
        target.unmerge_cells(str(merged))
    target.insert_rows(insert_row, amount)
    _shift_target_sheet_structures(target, insert_row, amount, moved_merges)
    for sheet, old_row, column, formula in formulas:
        destination_row = old_row + amount if sheet is target and old_row >= insert_row else old_row
        sheet.cell(destination_row, column).value = _rewrite_formula_for_insert(
            formula,
            formula_sheet=sheet.title,
            target_sheet=target.title,
            formula_row=old_row,
            insert_row=insert_row,
            amount=amount,
        )
    first_total_row = insert_row + amount
    for column in range(1, _logical_max_column(target, 3 if target.title.strip() == "Iteam表" else 4) + 1):
        cell = target.cell(first_total_row, column)
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.value = _extend_subtotal_to_previous_row(cell.value, first_total_row - 1)


def _copy_last_po_row(ws, target_row: int, *, source_row: int, header_row: int) -> None:
    source_dimension = ws.row_dimensions[source_row]
    if source_dimension.height:
        ws.row_dimensions[target_row].height = source_dimension.height
    for column in range(1, _logical_max_column(ws, header_row) + 1):
        source = ws.cell(source_row, column)
        destination = ws.cell(target_row, column)
        destination.value = None
        _copy_style(source, destination)
        if isinstance(source.value, str) and source.value.startswith("="):
            try:
                destination.value = Translator(
                    source.value,
                    origin=source.coordinate,
                ).translate_formula(destination.coordinate)
            except (TypeError, ValueError):
                destination.value = source.value
        font = copy(destination.font)
        font.color = "0000FF"
        destination.font = font


def _allocate_rows(
    workbook,
    data_workbook,
    ws,
    row_count: int,
    *,
    header_row: int,
) -> tuple[list[int], dict[str, list[int]], dict[str, int]]:
    columns = _header_columns(ws, header_row)
    total_row = _find_total_row(workbook, data_workbook, ws, header_row)
    last_row = _last_business_row(workbook, data_workbook, ws, header_row, total_row, columns)
    if last_row <= header_row:
        raise ValueError(f"360 排期工作表“{ws.title.strip()}”没有可仿照的历史 PO 行。")
    available = max(0, total_row - last_row - 1)
    inserted = max(0, row_count - available)
    if inserted:
        _insert_rows_preserving_workbook(workbook, ws, total_row, inserted)
    target_rows = list(range(last_row + 1, last_row + row_count + 1))
    for target_row in target_rows:
        _copy_last_po_row(ws, target_row, source_row=last_row, header_row=header_row)
    return target_rows, columns, {
        "header_row": header_row,
        "last_row": last_row,
        "total_row": total_row + inserted,
        "reused_blank_rows": min(row_count, available),
        "inserted_rows": inserted,
    }


def _excel_record_value(field: str, value: Any) -> Any:
    if value in (None, ""):
        return None
    if field in DATE_FIELDS:
        parsed = parse_iso_date(value)
        if parsed:
            return datetime(parsed.year, parsed.month, parsed.day)
    return value


def _set_field_format(cell, field: str) -> None:
    if field in DATE_FIELDS:
        cell.number_format = "yyyy-mm-dd"
    elif field in _IDENTIFIER_FIELDS:
        cell.number_format = "@"
    elif field in {"quantity", "outer_pack", "cartons"}:
        cell.number_format = "#,##0.##"
    elif field in {"unit_price_usd", "unit_price_hkd"}:
        cell.number_format = "0.0000"
    elif field in {"total_usd", "total_hkd"}:
        cell.number_format = "#,##0.00"


def _write_item_row(ws, row: int, record: dict[str, Any], columns: dict[str, list[int]]) -> None:
    for field, target_columns in columns.items():
        if field in _FORMULA_FIELDS:
            continue
        cell = ws.cell(row, target_columns[0], _excel_record_value(field, record.get(field)))
        _set_field_format(cell, field)

    quantity = columns.get("quantity", [None])[0]
    outer_pack = columns.get("outer_pack", [None])[0]
    cartons = columns.get("cartons", [None])[0]
    usd_columns = columns.get("unit_price_usd", [])
    hkd_columns = columns.get("unit_price_hkd", [])
    total_usd_columns = columns.get("total_usd", [])
    total_hkd_columns = columns.get("total_hkd", [])
    if quantity and outer_pack and cartons and not ws.cell(row, cartons).value:
        ws.cell(row, cartons).value = (
            f'=IFERROR({get_column_letter(quantity)}{row}/{get_column_letter(outer_pack)}{row},"")'
        )
    if usd_columns and hkd_columns and not ws.cell(row, hkd_columns[0]).value:
        ws.cell(row, hkd_columns[0]).value = f"={get_column_letter(usd_columns[0])}{row}*7.75"
    if quantity and usd_columns and total_usd_columns and not ws.cell(row, total_usd_columns[0]).value:
        ws.cell(row, total_usd_columns[0]).value = (
            f"={get_column_letter(usd_columns[0])}{row}*{get_column_letter(quantity)}{row}"
        )
    if quantity and hkd_columns and total_hkd_columns and not ws.cell(row, total_hkd_columns[-1]).value:
        ws.cell(row, total_hkd_columns[-1]).value = (
            f"={get_column_letter(hkd_columns[0])}{row}*{get_column_letter(quantity)}{row}"
        )


def _bind_linked_row_to_item(ws, row: int, *, item_sheet_name: str, item_row: int, header_row: int) -> None:
    for column in range(1, _logical_max_column(ws, header_row) + 1):
        cell = ws.cell(row, column)
        formula = cell.value
        if not isinstance(formula, str) or not formula.startswith("="):
            continue
        match = _SIMPLE_SHEET_REFERENCE.fullmatch(formula)
        if not match:
            continue
        referenced_sheet = (match.group("quoted") or match.group("plain") or "").replace("''", "'").strip()
        if referenced_sheet.casefold() != item_sheet_name.strip().casefold():
            continue
        prefix = f"'{item_sheet_name.replace(chr(39), chr(39) * 2)}'"
        cell.value = f"={prefix}!{match.group('column')}{item_row}"


def create_schedule_review_workbook(
    template: bytes,
    records: list[dict[str, Any]],
    output_path: Path,
    *,
    template_filename: str = "schedule.xlsx",
) -> Path:
    """Copy the complete Huaxing 360 schedule and append every PO to its three linked sheets."""
    if not template:
        raise ValueError("缺少当前360排期底表，无法生成回填副本。")
    prepared = [add_derived_fields(dict(record)) for record in records]
    workbook = load_complete_workbook_compatible(template, filename=template_filename)
    data_workbook = load_workbook_compatible(
        template,
        filename=template_filename,
        data_only=True,
    )
    try:
        item_sheet = _sheet_by_names(workbook, ("Iteam表", "ITEM表"))
        review_sheet = _sheet_by_names(workbook, ("正单评审表",))
        order_sheet = _sheet_by_names(workbook, ("接单表",))
        missing = [
            name
            for name, sheet in (
                ("Iteam表", item_sheet),
                ("正单评审表", review_sheet),
                ("接单表", order_sheet),
            )
            if sheet is None
        ]
        if missing:
            raise ValueError(f"华兴360排期缺少目标工作表：{'、'.join(missing)}")

        item_rows, item_columns, _ = _allocate_rows(
            workbook,
            data_workbook,
            item_sheet,
            len(prepared),
            header_row=3,
        )
        for row, record in zip(item_rows, prepared, strict=True):
            _write_item_row(item_sheet, row, record, item_columns)

        review_rows, _, _ = _allocate_rows(
            workbook,
            data_workbook,
            review_sheet,
            len(prepared),
            header_row=4,
        )
        for row, item_row in zip(review_rows, item_rows, strict=True):
            _bind_linked_row_to_item(
                review_sheet,
                row,
                item_sheet_name=item_sheet.title,
                item_row=item_row,
                header_row=4,
            )

        order_rows, _, _ = _allocate_rows(
            workbook,
            data_workbook,
            order_sheet,
            len(prepared),
            header_row=4,
        )
        for row, item_row in zip(order_rows, item_rows, strict=True):
            _bind_linked_row_to_item(
                order_sheet,
                row,
                item_sheet_name=item_sheet.title,
                item_row=item_row,
                header_row=4,
            )

        calculation = getattr(workbook, "calculation", None)
        if calculation is not None:
            calculation.fullCalcOnLoad = True
            calculation.forceFullCalc = True
            calculation.calcMode = "auto"
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        return output_path
    finally:
        data_workbook.close()
        workbook.close()


def create_summary_workbook(records: list[dict[str, Any]], output_path: Path) -> Path:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "360排期检查"
    ws.append([FIELD_TITLES[field] for field in EXPORT_FIELDS] + ["风险等级", "检查结果", "建议日期码"])
    for source in records:
        row = add_derived_fields(dict(source))
        ws.append([row.get(field) for field in EXPORT_FIELDS] + [row["risk_level"], "；".join(flag["text"] for flag in row["flags"]), row.get("expected_date_code")])
    _style_sheet(ws)
    summary = workbook.create_sheet("汇总")
    stats = summarize_records([add_derived_fields(dict(row)) for row in records])
    summary.append(["指标", "数值"])
    for label, field in (("订单行", "orders"), ("合同数", "contracts"), ("总数量", "quantity"), ("总箱数", "cartons"), ("高风险", "high_risk"), ("需跟进", "flagged")):
        summary.append([label, stats[field]])
    _style_sheet(summary)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
