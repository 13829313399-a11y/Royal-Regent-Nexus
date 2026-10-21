"""EDU-only grouped additions and explicit reservation consumption.

The caller allocates references and resolves the original ITEM reservation row.
This module never chooses between multiple material reservations or renumbers
historical orders. Original workbook cells move with Excel insertion semantics.
"""
from __future__ import annotations

from collections import defaultdict
from copy import copy
from datetime import date, datetime
import re
import xml.etree.ElementTree as ET

from openpyxl.comments import Comment
from openpyxl.chart.data_source import NumRef, StrRef, MultiLevelStrRef
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter, coordinate_to_tuple
from openpyxl.workbook.properties import CalcProperties

from .edu_schedule import normalize_edu_item


SHEETS = ('ITEM表', '接单表', '正单评审表')
PO_HEADERS = ('客出单日期', 'PO号码', 'Item#', '产品名称', '数量', '单价', '金额')


def _text(value):
    from app.services.customer_order_unified import _text as text
    return text(value)


def _fail(message):
    from app.services.customer_order_unified import CustomerOrderUnifiedError
    raise CustomerOrderUnifiedError(message)


def _columns(sheet):
    return (5, 8, 11, 6, 4) if sheet.title == 'ITEM表' else (4, 7, 10, 5, 3)


def _product(value):
    return normalize_edu_item(value).upper()


def _rgb(cell):
    color = cell.fill.fgColor
    if cell.fill.patternType != 'solid':
        return ''
    if color.type == 'rgb':
        return str(color.rgb)[-6:]
    if color.type == 'indexed' and color.indexed < len(COLOR_INDEX):
        return COLOR_INDEX[color.indexed][-6:]
    if color.type == 'theme' and cell.parent.parent.loaded_theme:
        theme_names = ('lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6', 'hlink', 'folHlink')
        try:
            root = ET.fromstring(cell.parent.parent.loaded_theme)
            element = root.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}' + theme_names[color.theme])
            color_element = list(element)[0]
            return color_element.attrib.get('lastClr') or color_element.attrib.get('val', '')
        except (ET.ParseError, IndexError, TypeError):
            return ''
    return ''


def _green(cell):
    rgb = _rgb(cell)
    if not re.fullmatch(r'[0-9a-fA-F]{6}', rgb):
        return False
    red, green, blue = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    # Pale formula cells have almost equal RGB channels and are not stock rows.
    return green - red >= 20 and green - blue >= 20


def is_reservation_row(sheet, row):
    """Recognize the evidenced green row, including blank/mistyped PO labels."""
    ref, product, quantity, _, _ = _columns(sheet)
    return bool(_text(sheet.cell(row, product).value) and
                _text(sheet.cell(row, ref).value) and
                all(_green(sheet.cell(row, col)) for col in (ref, product, quantity)))


def _boundary(sheet):
    ref, product, _, _, _ = _columns(sheet)
    markers = []
    for (r, c), cell in list(sheet._cells.items()):
        if r < 4 or c > 3 or cell.data_type == 'f':
            continue
        if _text(sheet.cell(r, ref).value) or _text(sheet.cell(r, product).value):
            continue
        if any(label in _text(cell.value) for label in ('取消单', '已走货', '已完成')):
            markers.append(r)
    return min(markers) if markers else max((r for (r, _), c in sheet._cells.items() if c.value is not None), default=3) + 1


def reservation_candidates(workbook, product):
    """Return original row coordinates; the caller must choose unambiguously."""
    if 'ITEM表' not in workbook.sheetnames:
        return []
    sheet = workbook['ITEM表']
    wanted = _product(product)
    return [dict(row=r, product=_text(sheet.cell(r, 8).value),
                 product_no=normalize_edu_item(sheet.cell(r, 8).value),
                 reference_no=_text(sheet.cell(r, 5).value),
                 po=_text(sheet.cell(r, 6).value), quantity=sheet.cell(r, 11).value)
            for r in range(4, _boundary(sheet))
            if _product(sheet.cell(r, 8).value) == wanted and is_reservation_row(sheet, r)]


def _insert(workbook, target, start, amount):
    """Use the shared insertion engine and protect its non-cell references locally."""
    from app.services.huakang_a_order_legacy.unified_schedule import _insert_rows
    from .new_order_excel import _rewrite_formula_for_insert

    def rewritten(value, sheet_name, row):
        if not isinstance(value, str) or not value:
            return value
        has_equals = value.startswith('=')
        result = _rewrite_formula_for_insert(value if has_equals else '=' + value,
            formula_sheet=sheet_name, target_sheet=target.title, formula_row=row,
            insert_row=start, amount=amount)
        return result if has_equals else result[1:]

    structures = []
    def chart_references(node, sheet_name, seen):
        if node is None or id(node) in seen:
            return
        seen.add(id(node))
        if isinstance(node, (NumRef, StrRef, MultiLevelStrRef)) and node.f:
            structures.append((node, 'f', rewritten(node.f, sheet_name, 1)))
        # Series defines its XML elements lazily during serialization, so its
        # populated descriptor attributes must also participate in this walk.
        for field in set(getattr(node, '__elements__', ())) | set(vars(node)):
            child = getattr(node, field, None)
            for value in child if isinstance(child, (tuple, list)) else [child]:
                if hasattr(value, '__elements__'):
                    chart_references(value, sheet_name, seen)
    for sheet in workbook:
        for chart in sheet._charts:
            chart_references(chart, sheet.title, set())
        for validation in sheet.data_validations.dataValidation:
            anchor = min((x.min_row for x in validation.sqref.ranges), default=1)
            for field in ('formula1', 'formula2'):
                value = getattr(validation, field)
                if value and not value.startswith('"'):
                    structures.append((validation, field, rewritten(value, sheet.title, anchor)))
        for region, rules in sheet.conditional_formatting._cf_rules.items():
            anchor = min((x.min_row for x in region.sqref.ranges), default=1)
            for rule in rules:
                if rule.formula:
                    structures.append((rule, 'formula', [rewritten(v, sheet.title, anchor) for v in rule.formula]))
        for name in sheet.defined_names.values():
            if name.attr_text:
                structures.append((name, 'attr_text', rewritten(name.attr_text, sheet.title, 1)))
    for name in workbook.defined_names.values():
        if name.attr_text:
            structures.append((name, 'attr_text', rewritten(name.attr_text, '', 1)))
    freeze = target.freeze_panes
    _insert_rows(workbook, target, start, amount)
    for structure, field, value in structures:
        setattr(structure, field, value)
    if freeze:
        r, c = coordinate_to_tuple(freeze)
        if r >= start:
            target.freeze_panes = f'{get_column_letter(c)}{r + amount}'
    for drawing in (*target._images, *target._charts):
        anchor = drawing.anchor
        if isinstance(anchor, str):
            r, c = coordinate_to_tuple(anchor)
            if r >= start:
                drawing.anchor = f'{get_column_letter(c)}{r + amount}'
        else:
            # The shared helper already moved _from; only a two-cell anchor's
            # far corner remains, otherwise the start would shift twice.
            for field in ('to',):
                marker = getattr(anchor, field, None)
                if marker is not None and marker.row >= start - 1:
                    marker.row += amount


def _slot(sheet, product):
    _, product_col, _, _, _ = _columns(sheet)
    boundary = _boundary(sheet)
    matches = [r for r in range(4, boundary) if _product(sheet.cell(r, product_col).value) == product]
    green = [r for r in matches if is_reservation_row(sheet, r)]
    return min(green) if green else max(matches) + 1 if matches else boundary


def _template(sheet, product, start):
    _, product_col, _, _, _ = _columns(sheet)
    ordinary = [r for r in range(4, _boundary(sheet)) if
                _text(sheet.cell(r, product_col).value) and not is_reservation_row(sheet, r)]
    same = [r for r in ordinary if _product(sheet.cell(r, product_col).value) == product]
    choices = same or ordinary
    return sheet.cell(min(choices, key=lambda r: abs(r - start)) if choices else max(4, start), 1)


def _copy_row(sheet, source, destination):
    from app.services.customer_order_unified import _copy_template_row
    _copy_template_row(sheet, source, destination)
    # A new ordinary order may have no ordinary source row in a stock-only sheet.
    for cell in sheet[destination]:
        if _green(cell):
            cell.fill = PatternFill()


def _identity(contract, po, product):
    from app.services.customer_order_huaxing_unified import edu_order_key
    return (*edu_order_key(contract, po), _product(product))


def _deduct(cell, quantity):
    from app.services.customer_order_unified import _number, _number_text
    value = cell.value
    if isinstance(value, str) and value.startswith('='):
        if value.startswith('={'):
            _fail(f'{cell.parent.title}!{cell.coordinate} 备料数量为数组公式，请先核对')
        expression = value[1:]
    else:
        number = _number(value)
        if number is None:
            _fail(f'{cell.parent.title}!{cell.coordinate} 备料数量为空或无效，不能当作零扣减')
        expression = _number_text(number)
    return '=' + expression + '-' + _number_text(quantity)


def _order_date(row):
    value = row.get('source_order_date') or row.get('received_date')
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10]) if value else None
    except ValueError:
        _fail('EDU 原收单日期无效，请核对日期后导出')


def _source_comment(cell, row):
    raw = _text(row.get('source_product_no'))
    if raw:
        cell.comment = Comment('原始产品型号：' + raw, '订单来源')


def _prepare(workbook, rows):
    from app.services.customer_order_unified import _number, _normalized_header
    if 'PO订单' in workbook.sheetnames and any(
            _normalized_header(workbook['PO订单'].cell(1, c).value) != _normalized_header(header)
            for c, header in enumerate(PO_HEADERS, 1)):
        _fail('EDU PO订单表头与已确认格式不一致，请核对后导出')
    item = workbook['ITEM表']
    existing = defaultdict(list)
    references = defaultdict(list)
    for r in range(4, item.max_row + 1):
        if not item.cell(r, 8).value or is_reservation_row(item, r):
            continue
        key = _identity(item.cell(r, 4).value, item.cell(r, 6).value, item.cell(r, 8).value)
        if any(key[:2]):
            existing[key].append(item.cell(r, 11).value)
        references[(_text(item.cell(r, 5).value), key[-1])].append(key)
    pending, seen, skipped = [], {}, 0
    for source in rows:
        if source.get('_edu_existing_order'):
            skipped += 1
            continue
        row = dict(source)
        row['product_no'] = normalize_edu_item(row.get('product_no'))
        quantity = _number(row.get('quantity'))
        reference = _text(row.get('reference_no'))
        key = _identity(row.get('contract_no'), row.get('po_no'), row.get('product_no'))
        if not reference or not key[-1] or not any(key[:2]) or quantity is None or quantity <= 0:
            _fail('EDU 缺少有效订单身份、产品或正数数量，不能导出')
        if key in existing or key in seen:
            quantities = existing.get(key, [seen.get(key)])
            if any(_number(q) != quantity for q in quantities):
                _fail('EDU 相同客户合同/PO及产品的数量冲突，不能重复扣备料')
            skipped += 1
            continue
        if references.get((reference, key[-1])):
            _fail('EDU 内部编号与产品已被另一客户订单占用，请重新预览分配')
        if 'PO订单' in workbook.sheetnames:
            po = workbook['PO订单']
            if any(_text(po.cell(r, 2).value) == reference and _product(po.cell(r, 3).value) == key[-1]
                   for r in range(2, po.max_row + 1)):
                _fail('EDU PO订单已存在相同内部编号及货号，但缺少一致正式单，请核对后再导出')
        selected = _text(row.get('edu_reservation_row'))
        candidates = reservation_candidates(workbook, row['product_no'])
        selected_row = None
        if selected:
            if not selected.isdigit() or int(selected) not in {r['row'] for r in candidates}:
                _fail('所选 EDU 备料行已失效或不属于当前货号，请重新预览并选择')
            selected_row = int(selected)
        elif candidates and not row.get('_edu_skip_reservation'):
            _fail('EDU 货号有备料行，请明确选择扣减行后导出')
        row['_quantity'] = quantity
        row['_reservation_cell'] = item.cell(selected_row, 11) if selected_row else None
        row['_source_date'] = _order_date(row)
        row['_summary_reservations'] = []
        if selected_row:
            _deduct(row['_reservation_cell'], quantity)  # Validate before any insertion.
            stock_reference = _text(item.cell(selected_row, 5).value)
            for name in SHEETS[1:]:
                sheet = workbook[name]
                for r in range(4, _boundary(sheet)):
                    if (is_reservation_row(sheet, r) and
                            _text(sheet.cell(r, 4).value) == stock_reference and
                            _product(sheet.cell(r, 7).value) == key[-1]):
                        cell = sheet.cell(r, 10)
                        if cell.data_type == 'f' and 'ITEM表!' in str(cell.value).replace("'", ''):
                            continue  # Already derives the ITEM reservation; never subtract twice.
                        if cell.data_type == 'f':
                            _fail(f'{name}!{cell.coordinate} 备料公式来源不明确，请先核对')
                        _deduct(cell, quantity)
                        row['_summary_reservations'].append(cell)
        seen[key] = quantity
        references[(reference, key[-1])].append(key)
        pending.append(row)
    return pending, skipped


def _append_po(workbook, rows):
    from app.services.customer_order_unified import _excel_value
    if 'PO订单' not in workbook.sheetnames:
        sheet = workbook.create_sheet('PO订单')
        for column, label in enumerate(PO_HEADERS, 1):
            sheet.cell(1, column, label)
    else:
        sheet = workbook['PO订单']
    start = max((r for (r, _), c in sheet._cells.items() if c.value is not None), default=1) + 1
    source = max((r for (r, c), cell in sheet._cells.items() if c == 2 and r >= 2 and cell.value), default=0)
    _insert(workbook, sheet, start, len(rows))
    for offset, row in enumerate(rows):
        r = start + offset
        if source:
            sheet.row_dimensions[r].height = sheet.row_dimensions[source].height
            for c in range(1, 8):
                sheet.cell(r, c)._style = copy(sheet.cell(source, c)._style)
        values = (row['_source_date'], row['reference_no'], row['product_no'],
                  row.get('product_name_zh') or row.get('product_name_en') or None,
                  _excel_value(row['_quantity']), None)
        for c, value in enumerate(values, 1):
            sheet.cell(r, c, value)
            if isinstance(value, str):
                sheet.cell(r, c).data_type = 's'
        sheet.cell(r, 1).number_format = 'yyyy-mm-dd'
        sheet.cell(r, 7, f'=IF(OR(E{r}="",F{r}=""),"",E{r}*F{r})')
        _source_comment(sheet.cell(r, 3), row)


def export_rows(workbook, rows):
    """Write each confirmed new detail once; return a JSON-safe audit summary."""
    from app.services import customer_order_unified as unified
    from app.services.customer_order_huaxing_unified import write_extras
    if not all(name in workbook.sheetnames for name in SHEETS):
        _fail('EDU 排期缺少 ITEM表、接单表或正单评审表')
    rows, skipped = _prepare(workbook, rows)
    if not rows:
        return {'written': 0, 'skipped_existing': skipped, 'reservation_deductions': []}
    grouped = defaultdict(list)
    for row in rows:
        grouped[_product(row['product_no'])].append(row)
    # Keep cell objects as coordinate handles: subsequent insertions move them.
    for name in SHEETS:
        sheet = workbook[name]
        for product, group in grouped.items():
            start = _slot(sheet, product)
            template_cell = _template(sheet, product, start)
            _insert(workbook, sheet, start, len(group))
            for offset, row in enumerate(group):
                r = start + offset
                _copy_row(sheet, template_cell.row, r)
                # Keys must be set now so later groups can find their own slots.
                ref, product_col, _, _, _ = _columns(sheet)
                sheet.cell(r, ref, row['reference_no'])
                sheet.cell(r, product_col, row['product_no'])
                row.setdefault('_cells', {})[name] = sheet.cell(r, 1)
    deductions = []
    for row in rows:
        item_row = row['_cells']['ITEM表'].row
        item = workbook['ITEM表']
        unified._write_item_row(item, item_row, row)
        write_extras(item, item_row, row)
        for (header_row, col), header in list(item._cells.items()):
            if header_row == 3 and _text(header.value) == '日期码':
                item.cell(item_row, col).value = row.get('date_code') or None
        # Business strings beginning '=' stay text; only our derived cells run.
        for col in [*range(1, 13), *range(14, 29), *[c for (r, c), cell in item._cells.items()
                    if r == 3 and _text(cell.value) in {'日期码', '备注', '联系人', 'PDQ'}]]:
            if isinstance(item.cell(item_row, col).value, str):
                item.cell(item_row, col).data_type = 's'
        _source_comment(item.cell(item_row, 8), row)
        order_row = row['_cells']['接单表'].row
        for name in SHEETS[1:]:
            sheet = workbook[name]
            r = row['_cells'][name].row
            unified._write_summary_row(sheet, r, row, item_row=item_row, item_sheet=item)
            for col in range(3, 8):
                if isinstance(sheet.cell(r, col).value, str):
                    sheet.cell(r, col).data_type = 's'
            _source_comment(sheet.cell(r, 7), row)
            if name == '接单表':
                sheet.cell(r, 14).value = unified._excel_value(unified._number(row.get('unit_price_hkd')))
                sheet.cell(r, 16).value = (f'=IF(OR(N{r}="",$P$1="",$P$1=0),"",N{r}/$P$1)'
                    if unified._number(sheet.cell(1, 16).value) else None)
            else:
                sheet.cell(r, 14).value = f'=IF(LEN(\'接单表\'!N{order_row})=0,"",\'接单表\'!N{order_row})'
                sheet.cell(r, 16).value = f'=IF(LEN(\'接单表\'!P{order_row})=0,"",\'接单表\'!P{order_row})'
            sheet.cell(r, 15).value = f'=IF(OR(J{r}="",N{r}=""),"",J{r}*N{r})'
            sheet.cell(r, 17).value = f'=IF(OR(J{r}="",P{r}=""),"",J{r}*P{r})'
        stock = row['_reservation_cell']
        if stock is not None:
            stock.value = _deduct(stock, row['_quantity'])
            deductions.append({'reference_no': row['reference_no'], 'product_no': row['product_no'],
                               'item_row': stock.row, 'quantity': unified._number_text(row['_quantity'])})
        for cell in row['_summary_reservations']:
            cell.value = _deduct(cell, row['_quantity'])
    _append_po(workbook, rows)
    workbook.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    return {'written': len(rows), 'skipped_existing': skipped, 'reservation_deductions': deductions}
