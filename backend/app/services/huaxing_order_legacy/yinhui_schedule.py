from __future__ import annotations

import io
import math
import re
import zipfile
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, BinaryIO, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.cell_range import CellRange

from .new_order_excel import (
    _copy_style,
    _extend_subtotal_to_previous_row,
    _rewrite_formula_for_insert,
    _shift_target_sheet_structures,
    load_complete_workbook_compatible,
    load_workbook_compatible,
)


EXCHANGE_RATE = 7.75

FIELD_TITLES = {
    "order_date": "出单日期",
    "so_no": "SO",
    "contract_no": "银辉合同号",
    "customer": "客名",
    "item_no": "产品编号",
    "product_name": "产品名称",
    "spec": "规格",
    "quantity": "PO数量",
    "case_pack": "装箱数量",
    "cartons": "总箱数",
    "manual": "说明书",
    "artwork": "彩盒",
    "customer_label": "客贴纸",
    "date_code": "日期码",
    "memo": "备注",
    "shipping_mark": "箱唛资料",
    "inspection_date": "验货日期",
    "factory_review_date": "华兴复期",
    "po_ship_date": "走货期",
    "unit_price_usd": "订单单价USD",
    "unit_price_hkd": "单价HK$",
    "total_hkd": "总金额HK$",
    "total_usd": "总金额USD",
    "factory_unit_price_hkd": "出厂价HK$",
    "factory_total_hkd": "出厂价总金额HK$",
}
EXPORT_FIELDS = list(FIELD_TITLES)


def normalize_label(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower().replace("（", "(").replace("）", ")")
    return re.sub(r"[\s/\\_()]+", "", text)


ALIASES = {
    normalize_label(title): field for field, title in FIELD_TITLES.items()
}
ALIASES.update({
    "来单日期": "order_date",
    "主合同号": "contract_no",
    "po号": "contract_no",
    "產品編號": "item_no",
    "產品型號": "item_no",
    "產品名稱": "product_name",
    "裝箱數量": "case_pack",
    "裝箱隻數": "case_pack",
    "装箱": "cartons",
    "外箱装箱数": "cartons",
    "箱数": "cartons",
    "验货期": "inspection_date",
    "走货日期": "po_ship_date",
    "po利宝": "customer_label",
    "单价": "unit_price_hkd",
    "单价hkd": "unit_price_hkd",
    "金额hk$": "total_hkd",
    "金额hkd": "total_hkd",
    "订单单价usd": "unit_price_usd",
    "单价hk$": "unit_price_hkd",
    "总金额hk$": "total_hkd",
    "总金额usd": "total_usd",
    "出厂价": "factory_unit_price_hkd",
    "出厂单价": "factory_unit_price_hkd",
})

DATE_FIELDS = {"order_date", "inspection_date", "factory_review_date", "po_ship_date"}
NUMBER_FIELDS = {
    "quantity", "case_pack", "cartons", "unit_price_usd", "unit_price_hkd",
    "total_hkd", "total_usd", "factory_unit_price_hkd", "factory_total_hkd",
}


@dataclass(frozen=True)
class SourceRef:
    kind: str
    path: Path
    member: str | None = None


def open_source_bytes(ref: SourceRef) -> bytes:
    if ref.kind == "zip":
        if not ref.member:
            raise ValueError("压缩包内文件名不能为空")
        with zipfile.ZipFile(ref.path) as archive:
            return archive.read(ref.member)
    return ref.path.read_bytes()


def source_bytes(source: str | Path | bytes | BinaryIO) -> bytes:
    if isinstance(source, bytes):
        return source
    if hasattr(source, "read"):
        return source.read()
    return Path(source).read_bytes()


def text_value(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def number_value(value: Any) -> float | int | None:
    if value in (None, "", "-") or isinstance(value, bool):
        return None
    try:
        number = float(str(value).replace(",", "").replace("$", "").strip())
    except ValueError:
        return None
    if not math.isfinite(number):
        return None
    return int(number) if number.is_integer() else round(number, 4)


def date_value(value: Any) -> str | None:
    if value in (None, "", "-"):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def as_date(value: Any) -> date | None:
    normalized = date_value(value)
    if not normalized:
        return None
    try:
        return date.fromisoformat(normalized[:10])
    except ValueError:
        return None


def decode_date_code(value: Any, reference_year: int | None = None) -> date | None:
    match = re.fullmatch(r"(\d{3})B(\d)RR", text_value(value) or "", re.I)
    if not match:
        return None
    ordinal, year_digit = int(match.group(1)), int(match.group(2))
    if ordinal < 1 or ordinal > 366:
        return None
    decade = ((reference_year or date.today().year) // 10) * 10
    year = decade + year_digit
    try:
        return date(year, 1, 1) + timedelta(days=ordinal - 1)
    except ValueError:
        return None


def make_date_code(value: Any) -> str | None:
    target = as_date(value)
    if not target:
        return None
    return f"{target.timetuple().tm_yday:03d}B{target.year % 10}RR"


def add_derived_fields(record: dict[str, Any], today: date | None = None) -> None:
    for field in DATE_FIELDS:
        record[field] = date_value(record.get(field))
    for field in NUMBER_FIELDS:
        record[field] = number_value(record.get(field))
    for field in set(FIELD_TITLES) - DATE_FIELDS - NUMBER_FIELDS:
        record[field] = text_value(record.get(field))

    quantity = record.get("quantity")
    pack = record.get("case_pack")
    usd = record.get("unit_price_usd")
    hkd = record.get("unit_price_hkd")
    factory = record.get("factory_unit_price_hkd")
    if record.get("cartons") is None and quantity is not None and pack:
        record["cartons"] = round(float(quantity) / float(pack), 2)
    if hkd is None and usd is not None:
        hkd = record["unit_price_hkd"] = round(float(usd) * EXCHANGE_RATE, 4)
    if record.get("total_usd") is None and quantity is not None and usd is not None:
        record["total_usd"] = round(float(quantity) * float(usd), 2)
    if record.get("total_hkd") is None and quantity is not None and hkd is not None:
        record["total_hkd"] = round(float(quantity) * float(hkd), 2)
    if record.get("factory_total_hkd") is None and quantity is not None and factory is not None:
        record["factory_total_hkd"] = round(float(quantity) * float(factory), 2)
    if not record.get("inspection_date") and record.get("po_ship_date"):
        ship = as_date(record["po_ship_date"])
        if ship:
            record["inspection_date"] = (ship - timedelta(days=5)).isoformat()

    ref = as_date(record.get("inspection_date")) or as_date(record.get("order_date"))
    decoded = decode_date_code(record.get("date_code"), ref.year if ref else None)
    record["date_code_date"] = decoded.isoformat() if decoded else None
    record["flags"] = build_flags(record, today=today)
    record["risk_level"] = (
        "high" if any(flag["level"] == "high" for flag in record["flags"])
        else "medium" if record["flags"] else "normal"
    )


def build_flags(record: dict[str, Any], today: date | None = None) -> list[dict[str, str]]:
    flags: list[dict[str, str]] = []
    required = {
        "so_no": "缺SO", "contract_no": "缺银辉合同号", "customer": "缺客名",
        "item_no": "缺产品编号", "product_name": "缺产品名称", "quantity": "缺PO数量",
        "case_pack": "缺装箱数量", "inspection_date": "缺验货日期",
        "po_ship_date": "缺走货期", "unit_price_usd": "缺订单单价USD",
    }
    for field, message in required.items():
        if record.get(field) in (None, "", 0):
            flags.append({"level": "high", "code": f"missing_{field}", "text": message})

    code = record.get("date_code")
    if not code:
        flags.append({"level": "medium", "code": "missing_date_code", "text": "日期码待填写"})
    elif not record.get("date_code_date"):
        flags.append({"level": "high", "code": "invalid_date_code", "text": "日期码格式应为日期序号+B年码+RR（例：183B6RR）"})

    quantity, pack, cartons = record.get("quantity"), record.get("case_pack"), record.get("cartons")
    if quantity is not None and pack and cartons is not None and abs(float(cartons) - float(quantity) / float(pack)) > 0.02:
        flags.append({"level": "high", "code": "carton_mismatch", "text": "总箱数与PO数量/装箱数量不一致"})

    ship = as_date(record.get("po_ship_date"))
    inspection = as_date(record.get("inspection_date"))
    if ship and inspection and inspection != ship - timedelta(days=5):
        flags.append({"level": "medium", "code": "inspection_review", "text": "验货日期不是走货期前5天，请复核"})
    reference = today or date.today()
    if ship and ship < reference:
        flags.append({"level": "high", "code": "overdue", "text": "走货期已过"})
    elif ship and (ship - reference).days <= 14:
        flags.append({"level": "medium", "code": "ship_soon", "text": "14天内走货"})

    for field, label in (("shipping_mark", "箱唛资料"), ("manual", "说明书"), ("artwork", "彩盒")):
        if not record.get(field):
            flags.append({"level": "medium", "code": f"missing_{field}", "text": f"{label}待确认"})
    return flags


def find_header(ws) -> tuple[int, dict[int, str]]:
    best_row, best_map = 0, {}
    for row in range(1, min(ws.max_row, 15) + 1):
        mapping: dict[int, str] = {}
        for col in range(1, min(ws.max_column, 100) + 1):
            field = ALIASES.get(normalize_label(ws.cell(row, col).value))
            if field:
                mapping[col] = field
        score = len(mapping) + (10 if {"item_no", "quantity", "customer"}.issubset(mapping.values()) else 0)
        best_score = len(best_map) + (10 if {"item_no", "quantity", "customer"}.issubset(best_map.values()) else 0)
        if score > best_score:
            best_row, best_map = row, mapping
    if len(best_map) < 6:
        raise ValueError(f"工作表“{ws.title}”未找到银辉排期表头")
    return best_row, best_map


def read_schedule(source: str | Path | bytes | BinaryIO, filename: str | None = None) -> dict[str, Any]:
    data = source_bytes(source)
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    candidates = sorted(workbook.worksheets, key=lambda ws: (ws.title.strip() not in {"Iteam表", "ITEM表"}, ws.title))
    last_error: Exception | None = None
    for ws in candidates:
        try:
            header_row, mapping = find_header(ws)
        except ValueError as exc:
            last_error = exc
            continue
        records: list[dict[str, Any]] = []
        for row in range(header_row + 1, ws.max_row + 1):
            record = {field: ws.cell(row, col).value for col, field in mapping.items()}
            identity = any(record.get(field) not in (None, "") for field in ("so_no", "contract_no", "item_no", "product_name"))
            if not identity:
                continue
            if text_value(record.get("item_no")) in {"合计", "总计", "小计"}:
                continue
            record["source_sheet"] = ws.title
            record["source_row"] = row
            add_derived_fields(record)
            records.append(record)
        if records:
            return {
                "filename": filename or "银辉排期.xlsx",
                "sheet": ws.title,
                "records": records,
                "summary": summarize_records(records),
                "meta": {"header_row": header_row, "columns": len(mapping), "sheet_names": workbook.sheetnames},
                "warnings": [],
            }
    raise ValueError(str(last_error or "Excel中未找到可识别的银辉排期明细"))


def summarize_records(records: Iterable[dict[str, Any]]) -> dict[str, Any]:
    rows = list(records)
    customers = Counter(str(row.get("customer") or "待确认") for row in rows)
    months: defaultdict[str, float] = defaultdict(float)
    for row in rows:
        ship = as_date(row.get("po_ship_date"))
        if ship:
            months[ship.strftime("%Y-%m")] += float(row.get("quantity") or 0)
    return {
        "orders": len(rows),
        "contracts": len({row.get("contract_no") for row in rows if row.get("contract_no")}),
        "quantity": round(sum(float(row.get("quantity") or 0) for row in rows), 2),
        "cartons": round(sum(float(row.get("cartons") or 0) for row in rows), 2),
        "total_usd": round(sum(float(row.get("total_usd") or 0) for row in rows), 2),
        "total_hkd": round(sum(float(row.get("total_hkd") or 0) for row in rows), 2),
        "high_risk": sum(row.get("risk_level") == "high" for row in rows),
        "flagged": sum(row.get("risk_level") != "normal" for row in rows),
        "top_customers": customers.most_common(8),
        "monthly_quantity": sorted(months.items()),
    }


_TOTAL_LABELS = {normalize_label(value) for value in ("合计", "總計", "总计", "小计", "小計")}
_SIMPLE_SHEET_REFERENCE = re.compile(
    r"^=\s*(?:'(?P<quoted>(?:[^']|'')+)'|(?P<plain>[^!]+))!"
    r"\$?(?P<column>[A-Z]{1,3})\$?(?P<row>\d+)\s*$"
)
_LINK_IDENTITY_FIELDS = ("so_no", "contract_no", "item_no", "product_name", "quantity")
_ITEM_COMPLETE_FIELDS = (
    "order_date", "so_no", "contract_no", "customer", "item_no", "product_name",
    "quantity", "unit_price_usd", "unit_price_hkd", "total_hkd", "total_usd",
    "factory_unit_price_hkd", "factory_total_hkd",
)
_IDENTIFIER_FIELDS = {"so_no", "contract_no", "item_no", "date_code"}


def _sheet_by_names(workbook, names: Iterable[str]):
    wanted = {str(name).strip().casefold() for name in names}
    return next(
        (sheet for sheet in workbook.worksheets if sheet.title.strip().casefold() in wanted),
        None,
    )


def _header_columns(ws, header_row: int) -> dict[str, list[int]]:
    columns: dict[str, list[int]] = defaultdict(list)
    for column in range(1, (ws.max_column or 1) + 1):
        field = ALIASES.get(normalize_label(ws.cell(header_row, column).value))
        if field:
            columns[field].append(column)
    return dict(columns)


def _effective_value(
    workbook,
    data_workbook,
    sheet_name: str,
    row: int,
    column: int,
    *,
    depth: int = 0,
) -> Any:
    formula_sheet = _sheet_by_names(workbook, (sheet_name,))
    data_sheet = _sheet_by_names(data_workbook, (sheet_name,))
    if formula_sheet is None:
        return None
    if data_sheet is not None:
        cached = data_sheet.cell(row, column).value
        if cached not in (None, ""):
            return cached
    raw = formula_sheet.cell(row, column).value
    if not isinstance(raw, str) or not raw.startswith("=") or depth >= 3:
        return raw
    match = _SIMPLE_SHEET_REFERENCE.fullmatch(raw)
    if not match:
        return None
    referenced_sheet = (match.group("quoted") or match.group("plain") or "").replace("''", "'").strip()
    return _effective_value(
        workbook,
        data_workbook,
        referenced_sheet,
        int(match.group("row")),
        column_index_from_string(match.group("column")),
        depth=depth + 1,
    )


def _find_total_row(workbook, data_workbook, ws, header_row: int) -> int:
    data_sheet = _sheet_by_names(data_workbook, (ws.title,))
    for row in range(header_row + 1, (ws.max_row or header_row) + 1):
        for column in range(1, (ws.max_column or 1) + 1):
            values = [ws.cell(row, column).value]
            if data_sheet is not None:
                values.append(data_sheet.cell(row, column).value)
            if any(
                isinstance(value, str)
                and not value.startswith("=")
                and normalize_label(value) in _TOTAL_LABELS
                for value in values
            ):
                return row
    return ws.max_row or header_row + 1


def _row_has_business_content(
    workbook,
    data_workbook,
    ws,
    row: int,
    columns: dict[str, list[int]],
) -> bool:
    for field in _LINK_IDENTITY_FIELDS:
        for column in columns.get(field, ()):
            value = _effective_value(workbook, data_workbook, ws.title, row, column)
            if value not in (None, "", "-") and normalize_label(value) not in _TOTAL_LABELS:
                return True
    return False


def _last_business_row(
    workbook,
    data_workbook,
    ws,
    header_row: int,
    total_row: int,
    columns: dict[str, list[int]],
) -> int:
    for row in range(total_row - 1, header_row, -1):
        if _row_has_business_content(workbook, data_workbook, ws, row, columns):
            return row
    return header_row


def _field_is_present(
    workbook,
    data_workbook,
    ws,
    row: int,
    columns: dict[str, list[int]],
    field: str,
) -> bool:
    return any(
        _effective_value(workbook, data_workbook, ws.title, row, column) not in (None, "", "-")
        for column in columns.get(field, ())
    )


def _complete_formula_row(
    workbook,
    data_workbook,
    ws,
    header_row: int,
    last_row: int,
    columns: dict[str, list[int]],
) -> int:
    required = _ITEM_COMPLETE_FIELDS if ws.title.strip().casefold() in {"iteam表", "item表"} else tuple(
        field for field in _LINK_IDENTITY_FIELDS if columns.get(field)
    )
    best_row = max(header_row, last_row)
    best_score = -1
    for row in range(last_row, header_row, -1):
        score = sum(
            _field_is_present(workbook, data_workbook, ws, row, columns, field)
            for field in required
        )
        if required and score == len(required):
            return row
        if score > best_score:
            best_row, best_score = row, score
    return best_row


def _insert_rows_preserving_workbook(workbook, target, insert_row: int, amount: int) -> None:
    if amount <= 0:
        return
    formulas = [
        (sheet, cell.row, cell.column, cell.value)
        for sheet in workbook.worksheets
        for cell in sheet._cells.values()
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    moved_merges = [
        CellRange(str(value))
        for value in target.merged_cells.ranges
        if value.max_row >= insert_row
    ]
    for merged in moved_merges:
        target.unmerge_cells(str(merged))
    target.insert_rows(insert_row, amount)
    _shift_target_sheet_structures(target, insert_row, amount, moved_merges)
    for sheet, old_row, column, formula in formulas:
        destination_row = old_row + amount if sheet is target and old_row >= insert_row else old_row
        sheet.cell(destination_row, column).value = _rewrite_formula_for_insert(
            formula,
            formula_sheet=sheet.title,
            target_sheet=target.title,
            formula_row=old_row,
            insert_row=insert_row,
            amount=amount,
        )
    first_total_row = insert_row + amount
    for row in range(first_total_row, (target.max_row or first_total_row) + 1):
        for column in range(1, (target.max_column or 1) + 1):
            cell = target.cell(row, column)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = _extend_subtotal_to_previous_row(cell.value, first_total_row - 1)


def _copy_template_row(
    ws,
    target_row: int,
    *,
    style_row: int,
    formula_row: int,
) -> None:
    source_dimension = ws.row_dimensions[style_row]
    if source_dimension.height:
        ws.row_dimensions[target_row].height = source_dimension.height
    for column in range(1, (ws.max_column or 1) + 1):
        style_source = ws.cell(style_row, column)
        formula_source = ws.cell(formula_row, column)
        destination = ws.cell(target_row, column)
        destination.value = None
        _copy_style(style_source, destination)
        if isinstance(formula_source.value, str) and formula_source.value.startswith("="):
            try:
                destination.value = Translator(
                    formula_source.value,
                    origin=formula_source.coordinate,
                ).translate_formula(destination.coordinate)
            except (TypeError, ValueError):
                destination.value = formula_source.value
        font = copy(destination.font)
        font.color = "0000FF"
        destination.font = font


def _allocate_rows(
    workbook,
    data_workbook,
    ws,
    row_count: int,
) -> tuple[list[int], dict[str, list[int]], dict[str, int]]:
    header_row, _ = find_header(ws)
    columns = _header_columns(ws, header_row)
    total_row = _find_total_row(workbook, data_workbook, ws, header_row)
    last_row = _last_business_row(workbook, data_workbook, ws, header_row, total_row, columns)
    style_row = max(header_row + 1, last_row)
    formula_row = _complete_formula_row(
        workbook,
        data_workbook,
        ws,
        header_row,
        last_row,
        columns,
    )
    available = max(0, total_row - last_row - 1)
    inserted = max(0, row_count - available)
    if inserted:
        _insert_rows_preserving_workbook(workbook, ws, total_row, inserted)
    target_rows = list(range(last_row + 1, last_row + row_count + 1))
    for target_row in target_rows:
        _copy_template_row(
            ws,
            target_row,
            style_row=style_row,
            formula_row=formula_row,
        )
    return target_rows, columns, {
        "header_row": header_row,
        "last_row": last_row,
        "style_row": style_row,
        "formula_row": formula_row,
        "total_row": total_row + inserted,
        "reused_blank_rows": min(row_count, available),
        "inserted_rows": inserted,
    }


def _excel_record_value(field: str, value: Any) -> Any:
    if value in (None, ""):
        return None
    if field in DATE_FIELDS:
        parsed = as_date(value)
        if parsed:
            return datetime(parsed.year, parsed.month, parsed.day)
    return value


def _preferred_column(columns: dict[str, list[int]], field: str) -> int | None:
    matches = columns.get(field, [])
    return matches[0] if matches else None


def _set_field_format(cell, field: str) -> None:
    if field in DATE_FIELDS:
        cell.number_format = "yyyy-mm-dd"
    elif field in _IDENTIFIER_FIELDS:
        cell.number_format = "@"
    elif field in {"quantity", "case_pack", "cartons"}:
        cell.number_format = "#,##0.##"
    elif field in {"unit_price_usd", "unit_price_hkd", "factory_unit_price_hkd"}:
        cell.number_format = "0.0000"
    elif field in {"total_usd", "total_hkd", "factory_total_hkd"}:
        cell.number_format = "#,##0.00"


def _write_item_row(ws, row: int, record: dict[str, Any], columns: dict[str, list[int]]) -> None:
    formula_fields = {
        "cartons", "total_hkd", "total_usd", "factory_unit_price_hkd", "factory_total_hkd",
    }
    for field, target_columns in columns.items():
        if field in formula_fields:
            continue
        selected_columns = target_columns[:1] if field == "memo" else target_columns
        value = _excel_record_value(field, record.get(field))
        for column in selected_columns:
            cell = ws.cell(row, column, value)
            _set_field_format(cell, field)

    quantity_column = _preferred_column(columns, "quantity")
    pack_column = _preferred_column(columns, "case_pack")
    cartons_column = _preferred_column(columns, "cartons")
    usd_column = _preferred_column(columns, "unit_price_usd")
    hkd_column = _preferred_column(columns, "unit_price_hkd")
    factory_unit_column = _preferred_column(columns, "factory_unit_price_hkd")
    if quantity_column and pack_column and cartons_column:
        fallback = record.get("cartons")
        fallback_text = str(float(fallback)) if fallback not in (None, "") else '""'
        ws.cell(row, cartons_column).value = (
            f"=IFERROR({get_column_letter(quantity_column)}{row}/"
            f"{get_column_letter(pack_column)}{row},{fallback_text})"
        )
        _set_field_format(ws.cell(row, cartons_column), "cartons")
    if quantity_column and hkd_column:
        for column in columns.get("total_hkd", ()):
            ws.cell(row, column).value = (
                f"={get_column_letter(quantity_column)}{row}*{get_column_letter(hkd_column)}{row}"
            )
            _set_field_format(ws.cell(row, column), "total_hkd")
    if quantity_column and usd_column:
        for column in columns.get("total_usd", ()):
            ws.cell(row, column).value = (
                f"={get_column_letter(quantity_column)}{row}*{get_column_letter(usd_column)}{row}"
            )
            _set_field_format(ws.cell(row, column), "total_usd")
    if factory_unit_column:
        factory_value = record.get("factory_unit_price_hkd")
        cell = ws.cell(row, factory_unit_column)
        if factory_value not in (None, ""):
            cell.value = factory_value
        elif usd_column:
            cell.value = f"={get_column_letter(usd_column)}{row}*{EXCHANGE_RATE}"
        _set_field_format(cell, "factory_unit_price_hkd")
    if quantity_column and factory_unit_column:
        for column in columns.get("factory_total_hkd", ()):
            cell = ws.cell(row, column)
            factory_total = record.get("factory_total_hkd")
            cell.value = factory_total if factory_total not in (None, "") else (
                f"={get_column_letter(factory_unit_column)}{row}*"
                f"{get_column_letter(quantity_column)}{row}"
            )
            _set_field_format(cell, "factory_total_hkd")


def _write_linked_row(
    ws,
    row: int,
    columns: dict[str, list[int]],
    *,
    item_sheet_name: str,
    item_row: int,
    item_columns: dict[str, list[int]],
) -> None:
    escaped_sheet = item_sheet_name.replace("'", "''")
    for field, target_columns in columns.items():
        source_column = _preferred_column(item_columns, field)
        if source_column is None:
            continue
        for column in target_columns:
            cell = ws.cell(row, column)
            cell.value = (
                f"='{escaped_sheet}'!{get_column_letter(source_column)}{item_row}"
            )
            _set_field_format(cell, field)


def _create_template_export(
    records: list[dict[str, Any]],
    output_path: Path,
    template_source: bytes | BinaryIO | str | Path,
    *,
    template_filename: str,
) -> Path:
    payload = source_bytes(template_source)
    workbook = load_complete_workbook_compatible(payload, filename=template_filename)
    data_workbook = load_workbook_compatible(payload, filename=template_filename, data_only=True)
    try:
        item_sheet = _sheet_by_names(workbook, ("Iteam表", "ITEM表"))
        review_sheet = _sheet_by_names(workbook, ("正单评审表",))
        order_sheet = _sheet_by_names(workbook, ("接单表",))
        if item_sheet is None or review_sheet is None or order_sheet is None:
            missing = [
                label
                for label, sheet in (
                    ("Iteam表", item_sheet),
                    ("正单评审表", review_sheet),
                    ("接单表", order_sheet),
                )
                if sheet is None
            ]
            raise ValueError(f"银辉排期缺少目标工作表：{'、'.join(missing)}")

        item_rows, item_columns, _ = _allocate_rows(
            workbook,
            data_workbook,
            item_sheet,
            len(records),
        )
        for row, record in zip(item_rows, records, strict=True):
            _write_item_row(item_sheet, row, record, item_columns)

        review_rows, review_columns, _ = _allocate_rows(
            workbook,
            data_workbook,
            review_sheet,
            len(records),
        )
        for row, item_row in zip(review_rows, item_rows, strict=True):
            _write_linked_row(
                review_sheet,
                row,
                review_columns,
                item_sheet_name=item_sheet.title,
                item_row=item_row,
                item_columns=item_columns,
            )

        order_rows, order_columns, _ = _allocate_rows(
            workbook,
            data_workbook,
            order_sheet,
            len(records),
        )
        for row, item_row in zip(order_rows, item_rows, strict=True):
            _write_linked_row(
                order_sheet,
                row,
                order_columns,
                item_sheet_name=item_sheet.title,
                item_row=item_row,
                item_columns=item_columns,
            )

        calculation = getattr(workbook, "calculation", None)
        if calculation is not None:
            calculation.fullCalcOnLoad = True
            calculation.forceFullCalc = True
            calculation.calcMode = "auto"
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        return output_path
    finally:
        data_workbook.close()
        workbook.close()


def create_export(
    records: list[dict[str, Any]],
    output_path: Path,
    template_source: bytes | BinaryIO | str | Path | None = None,
    *,
    template_filename: str = "schedule.xlsx",
) -> Path:
    if template_source is not None:
        prepared = []
        for source in records:
            record = dict(source)
            add_derived_fields(record)
            prepared.append(record)
        return _create_template_export(
            prepared,
            output_path,
            template_source,
            template_filename=template_filename,
        )
    wb = Workbook()
    summary = summarize_records(records)
    ws = wb.active
    ws.title = "汇总"
    ws.append(["指标", "数值"])
    for label, key in (("订单行", "orders"), ("合同数", "contracts"), ("PO数量", "quantity"),
                       ("总箱数", "cartons"), ("总金额USD", "total_usd"),
                       ("总金额HK$", "total_hkd"), ("需复核", "flagged"), ("高风险", "high_risk")):
        ws.append([label, summary[key]])
    detail = wb.create_sheet("Iteam明细")
    detail.append([FIELD_TITLES[field] for field in EXPORT_FIELDS] + ["系统校验"])
    for record in records:
        detail.append([record.get(field) for field in EXPORT_FIELDS] + ["；".join(flag["text"] for flag in record.get("flags", []))])
    header_fill = PatternFill("solid", fgColor="27364B")
    risk_fill = PatternFill("solid", fgColor="FCE8E6")
    for sheet in (ws, detail):
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        for col in range(1, sheet.max_column + 1):
            width = max(len(str(sheet.cell(row, col).value or "")) for row in range(1, sheet.max_row + 1))
            sheet.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 10), 32)
    for index, record in enumerate(records, start=2):
        if record.get("risk_level") == "high":
            for cell in detail[index]:
                cell.fill = risk_fill
    detail.auto_filter.ref = detail.dimensions
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
