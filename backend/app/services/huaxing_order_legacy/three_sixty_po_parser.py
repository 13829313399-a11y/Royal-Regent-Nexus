from __future__ import annotations

import io
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, BinaryIO, Iterable

import pdfplumber
from openpyxl import load_workbook

from .three_sixty_schedule import add_derived_fields, build_360_date_code


def _clean(value: Any) -> str | None:
    if value in (None, ""):
        return None
    text = re.sub(r"\s+", " ", str(value)).strip(" :\t\r\n")
    return text or None


def _date(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = re.sub(r"\s+\d{1,2}:\d{2}\s*(?:am|pm)?$", "", str(value).strip(), flags=re.I)
    for fmt in ("%d-%b-%Y", "%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text or None


def _number(value: Any) -> float | int | None:
    if value in (None, "") or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        match = re.search(r"-?[\d,.]+", str(value))
        if not match:
            return None
        number = float(match.group(0).replace(",", ""))
    return int(number) if number.is_integer() else number


def _search(pattern: str, text: str, flags: int = re.I | re.M) -> str | None:
    match = re.search(pattern, text, flags)
    return match.group(1).strip() if match else None


def _item_base(value: Any) -> str | None:
    digits = re.sub(r"\D", "", str(value or ""))
    return digits[:7] or None


def _normalise_safety(value: Any) -> str | None:
    text = _clean(value)
    if not text:
        return None
    upper = text.upper().replace("SAFETY WARNING", "SW")
    no_sw, no_qr = "NO SW" in upper, "NO QR" in upper
    if no_sw and no_qr:
        return "NO SW/QR"
    if no_sw:
        return "NO SW"
    if no_qr:
        return "NO QR"
    if "SW" in upper and "QR" in upper:
        # 正向要求可能包含 12L、品牌等有效版本资料，不能只缩写成
        # “+SW/QR”后丢掉 PO 原文。
        return text
    return text


def _product_version(im_value: Any, safety_value: Any, artwork: Any = None) -> str | None:
    im = _clean(im_value)
    safety = _normalise_safety(safety_value)
    parts: list[str] = []
    if im:
        parts.append(im if im.upper().startswith("IM") else f"IM : {im}")
    if safety:
        parts.append(safety)
    if not parts and artwork:
        parts.append(f"Artwork : {_clean(artwork)}")
    return "\n".join(parts) or None


def _release_title(text: str) -> tuple[str | None, int]:
    match = re.search(
        r"PURCHASE ORDER RELEASE[^\n]*?(RL-\d+-\d+)(?:\s*-\s*REVISION\s*(\d+))?",
        text,
        re.I,
    )
    if not match:
        return None, 0
    return match.group(1).upper(), int(match.group(2) or 0)


def read_pdf_text(source: str | Path | bytes | BinaryIO) -> tuple[str, int]:
    stream: Any = io.BytesIO(source) if isinstance(source, bytes) else source
    with pdfplumber.open(stream) as pdf:
        return (
            "\n".join(
                page.extract_text(x_tolerance=2, y_tolerance=3) or ""
                for page in pdf.pages
            ),
            len(pdf.pages),
        )


def _common(text: str, pages: int) -> dict[str, Any]:
    production_no, release_revision = _release_title(text)
    customer_po = _clean(_search(r"Customer PO Number\s*:\s*([^\n]+)", text))
    customer_release = _clean(_search(r"Customer Release No\.\s*:\s*([^\n]+)", text))
    if customer_release and customer_release.count("(") > customer_release.count(")"):
        customer_release += ")"
    our_contact = _clean(_search(r"Our Contact\s*:\s*([^\n]+)", text))
    return {
        "pages": pages,
        "production_no": production_no,
        "release_revision": release_revision,
        "revision_date": _date(_search(r"Revision Date\s*:\s*([^\n]+)", text)),
        "order_date": _date(_search(r"Order Date\s*:\s*([^\n]+)", text)),
        "customer_po": customer_po,
        "customer_po_type": _search(r"PURCHASE ORDER RELEASE\s*\(([^)]+)\)", text),
        "customer_release": customer_release,
        "our_contact": our_contact,
    }


def _release_customer(text: str, customer_release: Any) -> str:
    def without_contact_label(value: str) -> str:
        return re.split(r"\s+Contact Name\b.*", value, maxsplit=1, flags=re.I)[0].strip()

    customer = _clean(_search(r"End Customer\s*:\s*([^\n]+)", text))
    if customer:
        customer = without_contact_label(customer)
    if customer and customer.casefold() not in {"same as consignee", "n/a", "na"}:
        return customer
    release = _clean(customer_release) or ""
    # WPS Excel 的合并单元格有时会漏掉末尾右括号，仍按最后一组括号内容取客名。
    match = re.search(r"\(([^()]*)\)?\s*(?:FC[-\s].*)?$", release)
    if match:
        return match.group(1).strip()
    if "sales sample" in release.casefold():
        return "Sales Sample"
    ship_to = _clean(_search(r"Ship To \(Customer\)\s*:\s*([^\n]+)", text))
    if ship_to:
        ship_to = without_contact_label(ship_to)
    if ship_to and ship_to.casefold() not in {"same as consignee", "n/a", "na"} and "as per" not in ship_to.casefold():
        return ship_to
    warehouse = _clean(_search(r"Ship To \(Warehouse\)\s*:\s*([^\n]+)", text))
    if warehouse:
        warehouse = without_contact_label(warehouse)
    if warehouse and warehouse.casefold() not in {"n/a", "na"}:
        return warehouse
    return "MerchSource LLC"


def _release(text: str, meta: dict[str, Any]) -> list[dict[str, Any]]:
    production_no = meta.get("production_no")
    if not production_no:
        return []
    line = re.search(
        r"^\s*\d+\s+(\d{7,14}(?:\.\d+)?)\s+Each\s+([\d,]+(?:\.\d+)?)\s*$",
        text,
        re.I | re.M,
    )
    if not line:
        return []
    item_full, quantity_text = line.groups()
    following = text[line.end() : line.end() + 2600]
    item_section = re.split(
        r"\nThis Purchase Order Release|\n28/F,\s*Harbourside",
        following,
        maxsplit=1,
        flags=re.I,
    )[0]
    description = _clean(_search(r"^\s*([^\n]+)", item_section))
    artwork = _clean(
        _search(r"Artwork\s*:\s*(.+?)(?=\s+Brand Style Guide\s*:|$)", item_section)
    )
    # PDF/WPS 转换后可能把 “Safety” 拆成 “S afety”，并紧贴在 IM 内容后。
    safety_label = r"S\s*afety\s+Warning"
    im_value = _clean(
        _search(rf"IM\s*:\s*(.+?)(?={safety_label}\s*\+\s*QR\s*:|$)", item_section)
    )
    safety = _clean(_search(rf"{safety_label}\s*\+\s*QR\s*:\s*([^\n]+)", item_section))
    # 表格的左右两列被 PDF/WPS 文本层交错时，Safety 会插进多语言 IM 的中间；
    # 把 Safety 行之后、Plug Type 之前的语言续行接回 IM。
    im_continuation = _clean(
        _search(
            rf"{safety_label}\s*\+\s*QR\s*:[^\n]*\n([\s\S]*?)(?=\n?Plug Type\s*:)",
            item_section,
            flags=re.I | re.M,
        )
    )
    if im_value and im_continuation and not re.search(r"\w+\s*:", im_continuation):
        im_value = re.sub(r"\s*\+\s*", " + ", f"{im_value} {im_continuation}")
        im_value = re.sub(r"\s+", " ", im_value).strip()
    inspection = _date(
        _search(r"Planned Inspection Date\s*:\s*([0-9A-Za-z-]+)", text)
        or _search(r"Factory Commit Date\s*:\s*([0-9A-Za-z-]+)", text)
    )
    factory_no = _clean(_search(r"Factory No\s*:?\s*(\d{5})", text))
    customer_release = meta.get("customer_release")
    customer = _release_customer(text, customer_release)
    special_instruction = _clean(
        _search(
            r"SPECIAL INSTRUCTION\s*:\s*([\s\S]{0,1200}?)(?=Factory Commit Date History|$)",
            text,
        )
    )
    price_sticker = _clean(
        _search(r"(\$[\d.]+\s+price sticker[\s\S]{0,500}?smallest surface side)", text)
    )
    item_no = _item_base(item_full)
    version = _product_version(im_value, safety, artwork)
    row = {
        "order_date": meta.get("revision_date") or meta.get("order_date"),
        "po_no": meta.get("customer_po"),
        "customer_po": meta.get("customer_po"),
        "unified_customer_po": (
            'CINV-' + str(meta['customer_po'])
            if str(meta.get('customer_po_type') or '').upper() == 'CINV' and str(meta.get('customer_po') or '').isdigit()
            else meta.get('customer_po')
        ),
        "customer_release_no": customer_release,
        "production_no": production_no,
        "contract_no": _search(r"RL-(\d+)-\d+", production_no),
        "customer": customer,
        "item_no": item_no,
        "ship_item_no": item_no,
        "item_full": item_full,
        "product_name": description,
        "version": version,
        "version_source": (
            "PO IM / Safety Warning"
            if im_value or safety
            else ("PO Artwork（IM/Safety缺失时的暂用标识）" if artwork else None)
        ),
        "release_revision": meta.get("release_revision") or 0,
        "quantity": _number(quantity_text),
        "outer_pack": _number(_search(r"Master Carton Qty\s*:\s*([\d,.]+)", item_section)),
        "artwork": artwork,
        "packaging": _clean(_search(r"Package Type\s*:\s*([^\n]+)", item_section)),
        "vehicle": description,
        "inspection_date": inspection,
        "fcd_date": inspection,
        "date_code": build_360_date_code(inspection, factory_no),
        "merchandiser": meta.get("our_contact"),
        "special_remark": special_instruction,
        "customer_label": price_sticker,
        "third_party_inspection": (
            "需第三方验货" if re.search(r"Need\s+3rd\s+party\s+inspection", text, re.I) else None
        ),
        "barcode": _clean(_search(r"EAN Code\s*:\s*([^\n]+)", item_section)),
        "factory_no": factory_no,
        "container_type": _clean(
            _search(r"MS Container Type\s*\(FCL/LCL\)\s*:\s*([^\n]+)", text)
        ),
        "consolidated_id": _clean(_search(r"Consolidated ID\s*:\s*([^\n]+)", item_section)),
        "port_of_discharge": _clean(_search(r"Port of Discharge\s*:\s*([^\n]+)", text)),
        "transportation_mode": _clean(
            _search(r"Transportation Mode\s*:\s*([^\n]+)", text)
        ),
        "rl_order_type": _clean(_search(r"RL Order Type\s*:\s*([^\n]+)", text)),
        "source_sheet": "360 Release",
    }
    return [add_derived_fields(row)]


def _contract(text: str, meta: dict[str, Any]) -> list[dict[str, Any]]:
    contract = _search(r"PURCHASE ORDER[^\n]*?PO-(\d+)", text)
    rows: list[dict[str, Any]] = []
    pattern = re.compile(
        r"^\s*\d+\s+(\d{6,14})\s+Each\s+([\d,]+)\s+([\d,.]+)\s+([\d,.]+)\s*$",
        re.I | re.M,
    )
    for match in pattern.finditer(text):
        item_full, qty, unit, amount = match.groups()
        following = text[match.end() : match.end() + 500]
        item_no = _item_base(item_full)
        row = {
            "order_date": meta.get("order_date") or meta.get("revision_date"),
            "po_no": meta.get("customer_po"),
            "customer_po": meta.get("customer_po"),
            "contract_no": contract,
            "customer": "MerchSource LLC",
            "item_no": item_no,
            "ship_item_no": item_no,
            "item_full": item_full,
            "product_name": _clean(_search(r"^\s*([^\n]+)", following)),
            "quantity": _number(qty),
            "outer_pack": _number(
                _search(r"Master Carton Qty\s*:\s*([\d,.]+)", following)
            ),
            "unit_price_usd": _number(unit),
            "total_usd": _number(amount),
            "merchandiser": meta.get("our_contact"),
            "source_sheet": "360 Main PO",
        }
        rows.append(add_derived_fields(row))
    return rows


def parse_pdf_po(
    source: str | Path | bytes | BinaryIO,
    filename: str | None = None,
) -> dict[str, Any]:
    text, pages = read_pdf_text(source)
    if not text:
        return {
            "filename": filename or "po.pdf",
            "type": "pdf",
            "document_type": "unknown",
            "meta": {"pages": pages},
            "rows": [],
            "warnings": ["PDF未提取到文字，可能是扫描件。"],
            "text_preview": "",
        }
    release = bool(_release_title(text)[0]) or "ORDER RELEASE" in (filename or "").upper()
    meta = _common(text, pages)
    rows = _release(text, meta) if release else _contract(text, meta)
    meta.update({"text_chars": len(text), "row_count": len(rows)})
    warnings: list[str] = []
    if not rows:
        warnings.append("已读取PDF，但未识别到标准的360订单明细。")
    elif release and any(not row.get("version") for row in rows):
        warnings.append("PO未提供可识别的IM/Safety版本资料，版本列需人工确认。")
    elif release and any(str(row.get("version_source") or "").startswith("PO Artwork") for row in rows):
        warnings.append("部分 PO 没有 IM/Safety 资料，版本暂以 Artwork 标识，请跟单确认。")
    return {
        "filename": filename or "po.pdf",
        "type": "pdf",
        "document_type": "release" if release else "contract",
        "meta": meta,
        "rows": rows,
        "warnings": warnings,
        "text_preview": text[:1600],
    }


def _excel_rows(source: str | Path | bytes | BinaryIO) -> tuple[list[list[Any]], str]:
    raw: Any
    if isinstance(source, bytes):
        raw = io.BytesIO(source)
    elif hasattr(source, "read"):
        raw = io.BytesIO(source.read())
    else:
        raw = source
    workbook = load_workbook(raw, data_only=True, read_only=False)
    try:
        ws = workbook.active
        rows = [
            [ws.cell(row_no, col_no).value for col_no in range(1, ws.max_column + 1)]
            for row_no in range(1, ws.max_row + 1)
        ]
        return rows, ws.title
    finally:
        workbook.close()


def _excel_value_after(rows: list[list[Any]], label: str) -> str | None:
    target = re.sub(r"\s+", "", label).casefold()
    for row in rows:
        for index, value in enumerate(row):
            if value in (None, ""):
                continue
            text = str(value)
            compact = re.sub(r"\s+", "", text).casefold()
            if target not in compact:
                continue
            inline = re.search(re.escape(label) + r"\s*:\s*(.+)", text, re.I)
            if inline and _clean(inline.group(1)):
                return _clean(inline.group(1))
            for candidate in row[index + 1 :]:
                cleaned = _clean(candidate)
                if cleaned and cleaned != "0":
                    return _clean(re.sub(r"^:\s*", "", cleaned))
    return None


def _excel_as_release_text(rows: list[list[Any]]) -> str:
    title = next(
        (
            str(value)
            for row in rows
            for value in row
            if value not in (None, "") and "PURCHASE ORDER RELEASE" in str(value).upper()
        ),
        "",
    )
    item_block = ""
    quantity: Any = None
    for row in rows:
        for index, value in enumerate(row):
            text = str(value or "").strip()
            if re.match(r"^\d{7,14}(?:\.\d+)?(?:\n|$)", text):
                item_block = text
                numeric = [
                    candidate
                    for candidate in row[index + 1 :]
                    if _number(candidate) is not None and str(candidate).strip() not in {"0", "0.0"}
                ]
                quantity = numeric[-1] if numeric else None
                break
        if item_block:
            break
    item_lines = item_block.splitlines()
    item_full = item_lines[0].strip() if item_lines else ""
    body = "\n".join(item_lines[1:])
    parts = [title]
    for label in (
        "Revision Date",
        "Customer PO Number",
        "Customer Release No.",
        "Our Contact",
        "RL Order Type",
        "MS Container Type (FCL/LCL)",
    ):
        value = _excel_value_after(rows, label)
        if value:
            parts.append(f"{label} : {value}")
    if item_full:
        parts.append(f"1 {item_full} Each {quantity or ''}")
        parts.append(body)
    # Preserve shipping/factory/special-note rows for FCD, factory, customer and label parsing.
    for row in rows:
        values = [str(value) for value in row if value not in (None, "", "0", 0)]
        if values:
            parts.append(" ".join(values))
    return "\n".join(parts)


def parse_excel_po(
    source: str | Path | bytes | BinaryIO,
    filename: str | None = None,
) -> dict[str, Any]:
    rows, sheet_name = _excel_rows(source)
    text = _excel_as_release_text(rows)
    release = bool(_release_title(text)[0])
    meta = _common(text, len(re.findall(r"Page\s+\d+\s+of\s+\d+", text, re.I)) or 1)
    parsed_rows = _release(text, meta) if release else _contract(text, meta)
    meta.update({"sheet": sheet_name, "text_chars": len(text), "row_count": len(parsed_rows)})
    warnings: list[str] = []
    if not parsed_rows:
        warnings.append("Excel已读取，但未识别到标准的360订单明细。")
    elif release and any(not row.get("version") for row in parsed_rows):
        warnings.append("PO未提供可识别的IM/Safety版本资料，版本列需人工确认。")
    elif release and any(
        str(row.get("version_source") or "").startswith("PO Artwork")
        for row in parsed_rows
    ):
        warnings.append("部分 PO 没有 IM/Safety 资料，版本暂以 Artwork 标识，请跟单确认。")
    return {
        "filename": filename or "po.xlsx",
        "type": "excel",
        "document_type": "release" if release else "contract",
        "meta": meta,
        "rows": parsed_rows,
        "warnings": warnings,
        "text_preview": text[:1600],
    }


def parse_po_file(
    source: str | Path | bytes | BinaryIO,
    filename: str | None = None,
) -> dict[str, Any]:
    name = filename or (Path(source).name if isinstance(source, (str, Path)) else "upload")
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return parse_pdf_po(source, name)
    if suffix in {".xlsx", ".xlsm"}:
        return parse_excel_po(source, name)
    raise ValueError(f"不支持的文件类型：{suffix}")


def _completeness(row: dict[str, Any]) -> int:
    fields = (
        "production_no",
        "customer_po",
        "customer_release_no",
        "customer",
        "item_full",
        "product_name",
        "version",
        "quantity",
        "outer_pack",
        "inspection_date",
        "merchandiser",
        "barcode",
        "factory_no",
    )
    return sum(row.get(field) not in (None, "") for field in fields)


def merge_po_results(results: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    results = list(results)
    contracts = [
        row
        for result in results
        if result.get("document_type") == "contract"
        for row in result["rows"]
    ]
    releases = [
        row
        for result in results
        if result.get("document_type") == "release"
        for row in result["rows"]
    ]
    if not releases:
        return []

    selected: dict[str, dict[str, Any]] = {}
    for release in releases:
        # 当前 360 Release 样本每份只有一个明细行；同一 RL 的修订版可能会改货号。
        # 因此按 RL 选最高修订，不能把旧货号与新货号误当成两张新单。
        key = str(release.get("production_no") or "").upper()
        current = selected.get(key)
        candidate_rank = (
            int(release.get("release_revision") or 0),
            _completeness(release),
        )
        current_rank = (
            int(current.get("release_revision") or 0),
            _completeness(current),
        ) if current else (-1, -1)
        if current is None or candidate_rank > current_rank:
            selected[key] = dict(release)

    unique = list(selected.values())
    for release in unique:
        matches = [
            row
            for row in contracts
            if str(row.get("item_no")) == str(release.get("item_no"))
            and (
                not release.get("contract_no")
                or not row.get("contract_no")
                or str(row.get("contract_no")) == str(release.get("contract_no"))
            )
        ]
        special = " ".join(
            str(release.get(key) or "")
            for key in ("customer_release_no", "special_remark", "customer_label")
        ).lower()
        if matches and (
            str(release.get("customer_release_no") or "").lower().endswith("a")
            or "polybag" in special
        ):
            match = max(matches, key=lambda row: float(row.get("unit_price_usd") or 0))
        else:
            match = (
                min(matches, key=lambda row: float(row.get("unit_price_usd") or 0))
                if matches
                else None
            )
        if match:
            for field in ("unit_price_usd", "unit_price_hkd", "contract_no"):
                if not release.get(field):
                    release[field] = match.get(field)
        add_derived_fields(release)
    return unique
