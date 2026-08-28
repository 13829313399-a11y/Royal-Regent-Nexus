from __future__ import annotations

import os
import re
import shutil
from collections import Counter
from copy import copy
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import from_excel
from .new_order_excel import create_new_order_workbook


APP_DIR = Path(__file__).resolve().parent
DATA_DIR = APP_DIR / "data"
EXPORT_DIR = APP_DIR / "exports"
RUNTIME_UPLOAD_DIR = APP_DIR / "uploads"

CLIENT_NAMES = {"maxx": "Maxx", "shushupapa": "Shushupapa", "barter": "Barter"}
CLIENT_SHEETS = {
    "maxx": ("KFC公仔接单表", "接单表"),
    "shushupapa": ("接单表",),
    "barter": ("Iteam表", "接单表", "发票"),
}

ALIASES = {
    "status": ("系统",),
    "project": ("款号", "PROJECT", "PROJECTNO"),
    "po_date": ("来单日期", "客出单日期", "订单日期"),
    "customer": ("客户", "客名", "第三方客户"),
    "po_number": ("PONO", "采购订单NO", "正单合同号", "正单合同号S/CNO", "CONTRACT"),
    "customer_po": ("客PONO", "第三方客户/PONO", "第三方客户 PO NO#", "CUSTOMERPO"),
    "item": ("ITEMNO", "产品货号", "货号"),
    "product": ("产品名称", "名称"),
    "product_en": ("英文名称",),
    "quantity": ("每款总数量", "华兴产品数量", "产品数量", "数量"),
    "external_quantity": ("外厂产品数量",),
    "inner_pack": ("每款装箱个数", "华兴产品装箱个数", "产品装箱个数"),
    "outer_pack": ("外箱装箱数", "外箱总装箱数", "装箱数"),
    "cartons": ("总箱数",),
    "color_box": ("彩盒",),
    "manual": ("说明书",),
    "box_mark": ("箱唛资料", "箱唛资料(收到日期)"),
    "customer_label": ("客贴纸",),
    "label": ("贴纸",),
    "battery": ("电池",),
    "card": ("卡牌",),
    "country": ("走货国家/规格", "走货国家", "规格"),
    "unit_price": ("单价(HKD)", "单价HKD", "单价(USD)", "单价USD", "单价"),
    "amount": ("总金额(HKD)", "总金额HKD", "金额HKD", "总金额(USD)", "总金额USD", "金额USD"),
    "complete_date": ("完成日期",),
    "inspection_date": ("验货日期", "BV验货日期"),
    "ship_date": ("SHIPMENT", "走货期", "CRD"),
    "date_code": ("日期码",),
    "inspection_party": ("第三方公证行验货",),
    "inspection_result": ("验货结果",),
    "remarks": ("备注",),
}

BLUE_FILL = PatternFill("solid", fgColor="DCEBFF")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")


class EncryptedWorkbookError(ValueError):
    pass


def normalize_header(value: Any) -> str:
    text = str(value or "").upper().replace("：", ":")
    return re.sub(r"[\s._#:/\\()（）\-]+", "", text)


def normalize_key(value: Any) -> str:
    text = str(value or "").upper().strip()
    if text.endswith(".0"):
        text = text[:-2]
    return re.sub(r"\s+", "", text)


def normalize_item(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", normalize_key(value))


def display_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 1 < value < 100000:
        try:
            return from_excel(value).date()
        except (TypeError, ValueError, OverflowError):
            return None
    text = str(value or "").strip()[:10]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def is_encrypted(path: str | Path) -> bool:
    try:
        with open(path, "rb") as stream:
            signature = stream.read(8)
            if signature != bytes.fromhex("D0CF11E0A1B11AE1"):
                return False
        import olefile
        with olefile.OleFileIO(path) as ole:
            return ole.exists("EncryptedPackage")
    except OSError:
        return False


def active_path(client: str) -> Path:
    return RUNTIME_UPLOAD_DIR / "schedules" / client / "current_schedule.xlsx"


def schedule_path(client: str) -> Path | None:
    active = active_path(client)
    return active if active.exists() else None


def ensure_active(client: str) -> Path:
    active = active_path(client)
    if not active.exists():
        raise FileNotFoundError(f"当前没有 {CLIENT_NAMES[client]} 排期，请先由用户主动上传排期 .xlsx 文件。")
    return active


def _open(path: str | Path, *, data_only: bool = False):
    if is_encrypted(path):
        raise EncryptedWorkbookError("Excel 文件有打开密码，请先在 Excel 中解除密码后重新上传。")
    try:
        return load_workbook(path, data_only=data_only)
    except Exception as exc:
        if "not a zip" in str(exc).lower():
            raise ValueError("不是有效的 .xlsx 文件，或文件仍处于加密状态。") from exc
        raise


def _select_sheet(wb, client: str):
    preferred = CLIENT_SHEETS[client]
    for token in preferred:
        for name in wb.sheetnames:
            if token.lower() in name.strip().lower():
                return wb[name]
    return wb[wb.sheetnames[0]]


def _header_map(ws) -> tuple[int, dict[str, int], dict[int, str]]:
    alias_lookup = {}
    for field, aliases in ALIASES.items():
        for alias in aliases:
            alias_lookup[normalize_header(alias)] = field
    best: tuple[int, dict[str, int], dict[int, str]] | None = None
    for row_no in range(1, min(ws.max_row, 30) + 1):
        mapped: dict[str, int] = {}
        raw: dict[int, str] = {}
        for col in range(1, min(ws.max_column, 80) + 1):
            value = ws.cell(row_no, col).value
            if value in (None, ""):
                continue
            raw[col] = str(value).strip()
            normalized = normalize_header(value)
            field = alias_lookup.get(normalized)
            if field and field not in mapped:
                mapped[field] = col
        if len(mapped) >= 4 and (best is None or len(mapped) > len(best[1])):
            best = (row_no, mapped, raw)
    if not best:
        raise ValueError(f"工作表“{ws.title}”未找到可识别的排期表头。")
    return best


def inspect_schedule(path: str | Path, client: str) -> dict[str, Any]:
    wb = _open(path, data_only=False)
    try:
        ws = _select_sheet(wb, client)
        header_row, columns, raw = _header_map(ws)
        required = ["po_number", "item", "quantity"]
        missing = [name for name in required if name not in columns]
        if missing:
            raise ValueError("排期表缺少必要字段：" + "、".join(missing))
        return {"sheet": ws.title, "header_row": header_row, "columns": columns, "headers": raw}
    finally:
        wb.close()


def _record_flags(record: dict[str, Any], client: str, today: date) -> list[dict[str, str]]:
    flags: list[dict[str, str]] = []
    if not record.get("po_number"):
        flags.append({"level": "high", "code": "missing_po_number", "field": "po_number", "text": "缺订单号"})
    if not record.get("item"):
        flags.append({"level": "high", "code": "missing_item", "field": "item", "text": "缺货号"})
    if not record.get("ship_date"):
        flags.append({
            "level": "high" if client == "barter" else "medium",
            "code": "missing_ship_date",
            "field": "ship_date",
            "text": "缺走货期",
        })
    else:
        ship = parse_date(record["ship_date"])
        if ship and ship < today and str(record.get("inspection_result") or "").upper() not in {"PASS", "已走货"}:
            flags.append({"level": "high", "code": "past_ship_date", "field": "inspection_result", "text": "走货期已过"})
        elif ship and 0 <= (ship - today).days <= 14:
            flags.append({"level": "medium", "code": "near_ship_date", "field": "ship_date", "text": "14天内走货"})
    if client == "barter":
        required = (
            ("customer_po", "customer_po", "缺第三方客户 PO"),
            ("quantity", "quantity", "缺数量"),
            ("outer_pack", "outer_pack", "缺装箱数"),
            ("unit_price", "unit_price", "缺 HKD 单价"),
            ("amount", "amount", "缺 HKD 金额"),
            ("date_code", "date_code", "缺日期码"),
        )
        for key, field, message in required:
            if record.get(key) in (None, "", 0):
                flags.append({
                    "level": "high",
                    "code": f"missing_{key}",
                    "field": field,
                    "text": message,
                })
        if record.get("_ocr_used"):
            flags.append({
                "level": "medium",
                "code": "ocr_review",
                "field": "row",
                "text": "扫描合同已由 OCR 逐页识别，请对照原 PDF 复核关键字段",
            })
    if client == "shushupapa":
        item = normalize_key(record.get("item"))
        customer = normalize_key(record.get("customer"))
        product = str(record.get("product") or "")
        if "BJS" in customer and "带电" in product and item != "50002011":
            flags.append({"level": "high", "text": "BJs带电货号应为50002011"})
        if any(name in customer for name in ("WM", "TARGET", "COSTCO")) and "带电" in product and item != "50002010":
            flags.append({"level": "high", "text": "该客户带电货号应为50002010"})
        if "不带电" in product and item != "50002008":
            flags.append({"level": "high", "text": "不带电货号应为50002008"})
    return flags


def read_schedule(client: str, path: str | Path | None = None) -> dict[str, Any]:
    target = Path(path) if path else schedule_path(client)
    if target is None:
        return {
            "client": client, "client_name": CLIENT_NAMES[client], "path": "",
            "using_uploaded_copy": False, "empty": True, "sheet": "",
            "records": [],
            "summary": {"orders": 0, "quantity": 0, "amount": 0, "high_risk": 0, "flagged": 0},
        }
    info = inspect_schedule(target, client)
    wb = _open(target, data_only=True)
    try:
        ws = wb[info["sheet"]]
        columns = info["columns"]
        records = []
        for row_no in range(info["header_row"] + 1, ws.max_row + 1):
            record = {field: display_value(ws.cell(row_no, col).value) for field, col in columns.items()}
            if not any(record.get(key) not in (None, "") for key in ("po_number", "customer_po", "item", "quantity")):
                continue
            record["row"] = row_no
            record["flags"] = _record_flags(record, client, date.today())
            records.append(record)
        quantities = sum(float(r.get("quantity") or 0) for r in records if isinstance(r.get("quantity"), (int, float)))
        amounts = sum(float(r.get("amount") or 0) for r in records if isinstance(r.get("amount"), (int, float)))
        high = sum(any(f["level"] == "high" for f in r["flags"]) for r in records)
        flagged = sum(bool(r["flags"]) for r in records)
        return {
            "client": client, "client_name": CLIENT_NAMES[client], "path": str(target),
            "using_uploaded_copy": True, "empty": False, "sheet": info["sheet"],
            "records": records,
            "summary": {"orders": len(records), "quantity": quantities, "amount": round(amounts, 2), "high_risk": high, "flagged": flagged},
        }
    finally:
        wb.close()


def _copy_row_style(ws, source_row: int, target_row: int) -> None:
    for col in range(1, ws.max_column + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)
            target.number_format = source.number_format
            target.protection = copy(source.protection)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def _set(ws, row: int, columns: dict[str, int], field: str, value: Any, changes: list[str]) -> None:
    if field not in columns or value in (None, ""):
        return
    cell = ws.cell(row, columns[field])
    old = display_value(cell.value)
    if str(old or "").strip() == str(value or "").strip():
        return
    cell.value = value
    cell.fill = BLUE_FILL
    changes.append(f"{field}: {old or '空'} → {value}")


def _line_values(order: dict[str, Any], line: dict[str, Any]) -> dict[str, Any]:
    ship_value = line.get("ship_date") or order.get("ship_date")
    ship = parse_date(ship_value)
    return {
        "status": "已上系统",
        "project": order.get("project_no") or str(order.get("project") or "").split(" ", 1)[0],
        "po_date": order.get("po_date"), "customer": order.get("customer"),
        "po_number": line.get("po_number") or order.get("po_number"),
        "customer_po": line.get("customer_po") or order.get("customer_po"),
        "item": line.get("item_code"),
        "product": line.get("product_name_zh") or line.get("description"),
        "product_en": line.get("description"),
        "quantity": line.get("quantity"),
        "external_quantity": line.get("external_quantity"),
        "inner_pack": line.get("inner_pack"),
        "outer_pack": line.get("outer_pack"),
        "cartons": line.get("cartons"),
        "color_box": line.get("color_box"),
        "manual": line.get("manual"),
        "box_mark": line.get("box_mark"),
        "customer_label": line.get("customer_label"),
        "label": line.get("label"),
        "battery": line.get("battery"),
        "country": line.get("country"),
        "unit_price": line.get("unit_price"), "amount": line.get("amount"),
        "date_code": line.get("date_code"),
        # Maxx 规则文档明确：完成日期在 Shipment 前一周。
        "complete_date": (
            (ship - timedelta(days=7)).isoformat()
            if order.get("client") == "maxx" and ship
            else order.get("complete_date")
        ),
        # Barter 的 BV 验货日期按规则文档须"与客人商议确认"，留空待人工；
        # Maxx/Shushupapa 按文档为走货前一周。
        "inspection_date": (
            order.get("inspection_date") if order.get("client") == "barter"
            else (ship - timedelta(days=7)).isoformat() if ship else order.get("inspection_date")
        ),
        "ship_date": ship_value,
        "inspection_result": line.get("inspection_result") or order.get("inspection_result"),
        "remarks": line.get("shipment_note") or order.get("shipment_note"),
        "source_page": line.get("source_page"),
        "_ocr_used": bool(line.get("_ocr_used")),
        "_ocr_reconciled": bool(line.get("_ocr_reconciled")),
    }


def _read_order_history(client: str, path: str | Path) -> list[dict[str, Any]]:
    """Read unique PO rows from current and shipped-order sheets for deduplication."""
    wb = _open(path, data_only=True)
    try:
        records: list[dict[str, Any]] = []
        seen: set[tuple[str, ...]] = set()
        worksheets = list(wb.worksheets)
        if client == "barter":
            priority = {"Iteam表": 0, "接单表": 1, "发票": 2}
            worksheets.sort(key=lambda ws: priority.get(ws.title, 3))
        for ws in worksheets:
            try:
                header_row, columns, _ = _header_map(ws)
            except ValueError:
                continue
            if "po_number" not in columns:
                continue
            for row_no in range(header_row + 1, ws.max_row + 1):
                record = {
                    field: display_value(ws.cell(row_no, col).value)
                    for field, col in columns.items()
                }
                po_key = normalize_key(record.get("po_number"))
                if not po_key:
                    continue
                signature = (
                    po_key,
                    normalize_item(record.get("item")),
                    normalize_key(record.get("customer_po")),
                    normalize_key(record.get("quantity")),
                    normalize_key(record.get("ship_date")),
                )
                if signature in seen:
                    continue
                seen.add(signature)
                record["row"] = row_no
                record["sheet"] = ws.title
                records.append(record)
        return records
    finally:
        wb.close()


def apply_orders(client: str, orders: list[dict[str, Any]]) -> dict[str, Any]:
    target = ensure_active(client)
    EXPORT_DIR.mkdir(exist_ok=True)
    warnings: list[str] = []
    rows: list[dict[str, Any]] = []
    details: list[dict[str, Any]] = []
    existing_by_po: dict[str, list[dict[str, Any]]] = {}
    for record in _read_order_history(client, target):
        key = normalize_key(record.get("po_number"))
        if key:
            existing_by_po.setdefault(key, []).append(record)
    existing_orders = 0
    existing_lines_total = 0
    recognised_orders = 0
    for order in orders:
        if order.get("client") and order["client"] != client:
            warnings.append(f"{order.get('filename')}: 识别为 {order.get('client_name')}，已跳过。")
            continue
        recognised_orders += 1
        warnings.extend(order.get("warnings") or [])
        production_lines = [
            line
            for line in (order.get("lines") or [])
            if not line.get("is_charge")
        ]
        existing_in_order = 0
        new_in_order = 0
        po_quantity = 0.0
        schedule_quantity = 0.0
        for line in production_lines:
            line_po = line.get("po_number") or order.get("po_number")
            po_key = normalize_key(line_po)
            existing_records = existing_by_po.get(po_key, []) if po_key else []
            if existing_records:
                existing_in_order += 1
                existing_lines_total += 1
                line_po_quantity = float(line.get("quantity") or 0)
                line_schedule_quantity = sum(
                    float(record.get("quantity") or 0)
                    for record in existing_records
                    if isinstance(record.get("quantity"), (int, float))
                )
                po_quantity += line_po_quantity
                schedule_quantity += line_schedule_quantity
                details.append({
                    "type": "existing",
                    "po": line_po,
                    "schedule_rows": len(existing_records),
                    "po_quantity": line_po_quantity,
                    "schedule_quantity": line_schedule_quantity,
                    "schedule_sheets": sorted({
                        str(record.get("sheet") or "") for record in existing_records
                    }),
                })
                continue
            rows.append(_line_values(order, line))
            new_in_order += 1
        if existing_in_order:
            existing_orders += 1
            quantity_note = (
                f"，数量合计一致 {schedule_quantity:g}"
                if po_quantity and abs(schedule_quantity - po_quantity) < 0.0001
                else f"，PO数量 {po_quantity:g} / 排期数量 {schedule_quantity:g}，疑似修改单"
            )
            warnings.append(
                f"{order.get('po_number')}: {existing_in_order}/{len(production_lines)} 个分批已在当前或已出货排期"
                f"{quantity_note}；已去重"
                + (f"，仅生成其余 {new_in_order} 个新分批。" if new_in_order else "，未重复生成。")
            )
        for line in order.get("lines") or []:
            if line.get("is_charge"):
                warnings.append(
                    f"{order.get('po_number')} / {line.get('description')} 为工模或费用项，未写入生产排期。"
                )
    if not rows:
        if existing_orders:
            return {
                "ok": True,
                "modified": 0,
                "added": 0,
                "details": details,
                "rows": [],
                "warnings": list(dict.fromkeys(warnings)),
                "export_file": None,
                "mode": "new_order_only",
                "meta": {
                    "success_count": recognised_orders,
                    "failed_count": 0,
                    "ignored_count": 0,
                    "existing_count": existing_orders,
                    "existing_line_count": existing_lines_total,
                },
                "message": (
                    f"已识别 {recognised_orders} 份 PO；其中 {existing_orders} 份已在当前排期，"
                    "没有新单，系统未生成空白 Excel。"
                ),
            }
        raise ValueError(
            "PO 文件已读取，但没有识别到可写入的生产产品行；系统已拦截，未生成空白 Excel。"
        )
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    export_name = f"{CLIENT_NAMES[client]}_新单_{stamp}.xlsx"
    create_info = create_new_order_workbook(
        target,
        EXPORT_DIR / export_name,
        rows,
        ALIASES,
        sheet_names=CLIENT_SHEETS[client],
        sheet_title="新单",
    )
    details.extend(
        {"type": "new", "item": row.get("item"), "po": row.get("po_number")}
        for row in rows
    )
    return {
        "ok": True,
        "modified": 0,
        "added": len(rows),
        "details": details,
        "rows": rows,
        "warnings": list(dict.fromkeys(warnings)),
        "export_file": export_name,
        "export_info": create_info,
        "mode": "new_order_only",
        "meta": {
            "success_count": recognised_orders,
            "failed_count": 0,
            "ignored_count": 0,
            "existing_count": existing_orders,
            "existing_line_count": existing_lines_total,
        },
        "message": (
            f"已按新单生成 {len(rows)} 行 Excel，原排期未修改。"
            + (f"另有 {existing_orders} 份 PO 已在当前排期，未重复生成。" if existing_orders else "")
        ),
    }


def create_summary(client: str) -> str:
    data = read_schedule(client)
    wb = Workbook()
    ws = wb.active
    ws.title = "排期汇总"
    fields = ["status", "po_date", "customer", "po_number", "customer_po", "item", "product", "quantity", "cartons", "unit_price", "amount", "inspection_date", "ship_date", "date_code", "inspection_result", "remarks"]
    labels = ["系统", "来单日期", "客户", "订单号", "客PO", "货号", "产品名称", "数量", "总箱数", "单价USD", "金额USD", "验货日期", "走货期", "日期码", "验货结果", "备注"]
    for col, label in enumerate(labels, 1):
        cell = ws.cell(1, col, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17324D")
    for row_no, record in enumerate(data["records"], 2):
        for col, field in enumerate(fields, 1):
            ws.cell(row_no, col, record.get(field, ""))
        if record["flags"]:
            for cell in ws[row_no]:
                cell.fill = YELLOW_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [12, 12, 14, 18, 18, 16, 30, 12, 12, 12, 14, 12, 12, 14, 12, 36]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + index)].width = width
    EXPORT_DIR.mkdir(exist_ok=True)
    name = f"{CLIENT_NAMES[client]}_排期汇总_{datetime.now().strftime('%Y%m%d-%H%M%S-%f')}.xlsx"
    wb.save(EXPORT_DIR / name)
    wb.close()
    return name
