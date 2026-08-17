from __future__ import annotations

import io
import re
from collections import defaultdict
from copy import copy
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, BinaryIO, Iterable, Mapping, Sequence

import pdfplumber
from openpyxl.formula.translate import Translator
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.cell_range import CellRange

from . import new_order_excel


ITEM_SHEET = "ITEM表"
REVIEW_SHEET = "正单评审表"
ORDER_SHEET = "接单表"
EXCHANGE_RATE = Decimal("7.75")

CUSTOMERS = {
    "DLR": ("DLR（美国乐园）", "美国乐园"),
    "WDW": ("WDW（美国乐园）", "美国乐园"),
    "TDSE": ("TDSE（欧洲）", "欧洲"),
    "LIVERPOOL": ("LIVERPOOL（利物浦）", "墨西哥/利物浦"),
    "INTERNATIONALTOYWORLD": ("菲律宾（INTERNATIONALTOYWORLD）", "菲律宾"),
    "ALSHAYA": ("ALshaya（阿拉伯）", "阿拉伯/中东"),
    "HYUNDAI": ("HYUNDAI（韩国现代商超）", "韩国"),
    "JAPAN": ("JApan（日本）", "日本"),
}


def _clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _key(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", _clean(value).upper())


def _label_key(value: Any) -> str:
    return re.sub(r"[\s\u3000_／/()（）【】\[\]：:·.,，。-]+", "", _clean(value)).casefold()


def _number(value: Any) -> int | float | None:
    text = re.sub(r"[^\d.\-]", "", str(value or ""))
    if text in {"", "-", "."}:
        return None
    try:
        result = Decimal(text)
    except InvalidOperation:
        return None
    if result == result.to_integral_value():
        return int(result)
    return float(result)


def _iso_date(value: Any) -> str | None:
    text = _clean(value)
    if not text:
        return None
    for fmt in (
        "%m/%d/%Y", "%m/%d/%y", "%d-%b-%Y", "%d-%b-%y",
        "%Y-%m-%d", "%Y/%m/%d",
    ):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _minus_days(value: Any, days: int) -> str | None:
    normalized = _iso_date(value) or _clean(value)
    try:
        return (date.fromisoformat(normalized) - timedelta(days=days)).isoformat()
    except ValueError:
        return None


def _schedule_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        try:
            return from_excel(value).date().isoformat()
        except (TypeError, ValueError, OverflowError):
            return ""
    return _iso_date(value) or ""


def _revision(filename: str, text: str = "") -> int:
    candidates = [
        int(value)
        for value in re.findall(
            r"(?:\bR|REVISION\s+)(?:EVISION\s+)?0*(\d+)\b",
            f"{filename}\n{text[:1200]}",
            re.I,
        )
    ]
    return max(candidates, default=0)


def _pdf_text(source: str | Path | bytes | BinaryIO) -> tuple[str, int]:
    if isinstance(source, bytes):
        payload: str | Path | BinaryIO = io.BytesIO(source)
    elif isinstance(source, (str, Path)):
        payload = str(source)
    else:
        payload = source
        try:
            source.seek(0)
        except (AttributeError, OSError):
            pass
    with pdfplumber.open(payload) as pdf:
        pages = [page.extract_text(x_tolerance=2, y_tolerance=3) or "" for page in pdf.pages]
    return "\n".join(pages).strip(), len(pages)


def _record(
    *, po_number: str, customer_code: str, item: str, description: str,
    quantity: Any, case_pack: Any, ship_date: str | None, unit_price_usd: Any,
    source_type: str,
) -> dict[str, Any]:
    customer, country = CUSTOMERS[customer_code]
    return {
        "po_number": po_number,
        "customer_code": customer_code,
        "customer": customer,
        "country": country,
        "item": item,
        "description": _clean(description),
        "quantity": quantity,
        "case_pack": case_pack,
        "ship_date": ship_date,
        "unit_price_usd": unit_price_usd,
        "source_type": source_type,
    }


def _parse_theme_park(text: str) -> list[dict[str, Any]]:
    po_match = re.search(r"\b([LW]-\d{6,})\b", text)
    if not po_match:
        return []
    po_number = po_match.group(1)
    customer_code = "DLR" if po_number.startswith("L-") else "WDW"
    date_section = re.search(
        r"EARLIEST\s+SHIP\s+DATE.*?REVISION\s+NUMBER",
        text,
        re.I | re.S,
    )
    dates = re.findall(r"\b\d{1,2}/\d{1,2}/\d{4}\b", date_section.group(0) if date_section else text[:1600])
    ship_date = _iso_date(dates[0]) if dates else None
    line_pattern = re.compile(
        r"(?m)^(\d{10})\s+(.+?)\s+([\d,]+)\s+EA\s+\d+\s*/\s*EA\s+([\d.]+)\s+",
        re.I,
    )
    matches = list(line_pattern.finditer(text))
    rows: list[dict[str, Any]] = []
    for index, match in enumerate(matches):
        segment = text[match.start(): matches[index + 1].start() if index + 1 < len(matches) else len(text)]
        case_match = re.search(r"CASE\s*/\s*PK\s*:\s*([\d,]+)", segment, re.I)
        rows.append(_record(
            po_number=po_number,
            customer_code=customer_code,
            item=match.group(1),
            description=match.group(2),
            quantity=_number(match.group(3)),
            case_pack=_number(case_match.group(1)) if case_match else None,
            ship_date=ship_date,
            unit_price_usd=_number(match.group(4)),
            source_type="theme_park",
        ))
    return rows


def _parse_tdse(text: str) -> list[dict[str, Any]]:
    po_match = re.search(r"\b(W\d{4,})\b", text)
    if not po_match:
        return []
    section = re.search(r"Order\s+Ship\s+Anticipate\s+Cancel\s+Date(.*?)Ship\s+via", text, re.I | re.S)
    date_text = section.group(1) if section else text[:1800]
    split_dates = re.search(
        r"(\d{1,2}-[A-Z]{3})-\s+(\d{1,2}-[A-Z]{3})-\s+(\d{1,2}-[A-Z]{3})-",
        date_text,
        re.I,
    )
    year_match = re.search(r"20\d{2}", date_text[split_dates.end():] if split_dates else date_text)
    ship_date = (
        _iso_date(f"{split_dates.group(2)}-{year_match.group(0)}")
        if split_dates and year_match
        else None
    )
    style_matches = list(re.finditer(r"(?m)^\d+\s+(\d{10})\b([^\n]*)$", text))
    rows: list[dict[str, Any]] = []
    for index, style_match in enumerate(style_matches):
        segment = text[style_match.start(): style_matches[index + 1].start() if index + 1 < len(style_matches) else len(text)]
        pack_match = re.search(r"\b(\d+)\s*/\s*(\d+)\s*$", style_match.group(2), re.M)
        amount_matches = re.findall(r"(?m)^.*?\b([\d,]+)\s+([\d]+\.\d{2,4})\s*$", segment)
        quantity, unit_price = (amount_matches[-1] if amount_matches else (None, None))
        rows.append(_record(
            po_number=po_match.group(1),
            customer_code="TDSE",
            item=style_match.group(1),
            description="",
            quantity=_number(quantity),
            case_pack=_number(pack_match.group(1)) if pack_match else None,
            ship_date=ship_date,
            unit_price_usd=_number(unit_price),
            source_type="tdse",
        ))
    return rows


def _parse_f_order(text: str) -> list[dict[str, Any]]:
    po_match = re.search(r"\b(F\d{10,})\b", text)
    if not po_match:
        return []
    classification_match = re.search(
        r"(?m)^\s*(LIVERPOOL|INTERNATIONALTOYWORLD|ALSHAYA|HYUNDAI)\s*:",
        text,
        re.I,
    )
    if not classification_match:
        return []
    customer_code = classification_match.group(1).upper()
    date_section = re.search(r"ORDERED\s+SHIP\s+ON\s+ANTICIPATE\s+CANCEL\s+AFTER.*?\n([^\n]+)", text, re.I)
    dates = re.findall(r"\b\d{2}/\d{2}/\d{2}\b", date_section.group(1) if date_section else text[:1200])
    ship_date = _iso_date(dates[1]) if len(dates) > 1 else None
    pattern = re.compile(
        r"(?m)^(\d{10})\s+(.+?)\s+Multi\s+([\d.]+)\s+(?:NO\s+SIZE|Not\s+Sized)\s+([\d,]+)\s*$",
        re.I,
    )
    rows: list[dict[str, Any]] = []
    for match in pattern.finditer(text):
        segment = text[match.start():]
        case_match = re.search(r"Case\s+Pack\s*=\s*([\d,]+)", segment, re.I)
        rows.append(_record(
            po_number=po_match.group(1),
            customer_code=customer_code,
            item=match.group(1),
            description=match.group(2),
            quantity=_number(match.group(4)),
            case_pack=_number(case_match.group(1)) if case_match else None,
            ship_date=ship_date,
            unit_price_usd=_number(match.group(3)),
            source_type="international",
        ))
    return rows


def _parse_japan(text: str) -> list[dict[str, Any]]:
    po_match = re.search(r"\b(V\d{4,})\b", text)
    if not po_match:
        return []
    header = text[:1400]
    dates = re.findall(r"\b\d{1,2}/\d{1,2}/\d{2}\b", header)
    ship_date = _iso_date(dates[1]) if len(dates) > 1 else None
    pattern = re.compile(
        r"(?m)^(\d{10})\s+(.+?)\s+NO\s+COLOR\s*\n[^\n]*?\s+([\d.]+)\s+NO\s+SIZE\s+([\d,]+)\s*$",
        re.I,
    )
    rows: list[dict[str, Any]] = []
    for match in pattern.finditer(text):
        segment = text[match.start():]
        case_match = re.search(r"CASE\s+PACK\s*=\s*([\d,]+)", segment, re.I)
        rows.append(_record(
            po_number=po_match.group(1),
            customer_code="JAPAN",
            item=match.group(1),
            description=match.group(2),
            quantity=_number(match.group(4)),
            case_pack=_number(case_match.group(1)) if case_match else None,
            ship_date=ship_date,
            unit_price_usd=_number(match.group(3)),
            source_type="japan",
        ))
    return rows


def parse_text(text: str, filename: str) -> dict[str, Any]:
    normalized_name = Path(filename).stem
    if re.search(r"(?:^|[ _-])TL(?:$|[ _-]|\s*\(R\d+\))", normalized_name, re.I):
        return {
            "filename": filename,
            "po_number": "",
            "revision": _revision(filename, text),
            "rows": [],
            "warnings": ["识别为 TL/条款附件，未作为订单明细导入。"],
            "ignored": True,
        }
    parsers = (
        _parse_theme_park,
        _parse_tdse,
        _parse_f_order,
        _parse_japan,
    )
    rows = next((parsed for parser in parsers if (parsed := parser(text))), [])
    po_number = _clean(rows[0].get("po_number")) if rows else ""
    warnings: list[str] = []
    if not rows:
        warnings.append("未识别到受支持的迪士尼 PO 明细；支持 DLR、WDW、TDSE、国际 F 单及日本 V 单。")
    return {
        "filename": filename,
        "po_number": po_number,
        "revision": _revision(filename, text),
        "rows": rows,
        "warnings": warnings,
        "ignored": False,
    }


def parse_po(source: str | Path | bytes | BinaryIO, filename: str) -> dict[str, Any]:
    text, pages = _pdf_text(source)
    if not text:
        raise ValueError("迪士尼 PO 没有可读取的文字层")
    result = parse_text(text, filename)
    result["pages"] = pages
    return result


def merge_revisions(documents: Iterable[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[str]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    ignored: list[dict[str, Any]] = []
    for document in documents:
        po_number = _key(document.get("po_number"))
        if po_number:
            grouped[po_number].append(document)
        else:
            ignored.append(document)
    selected: list[dict[str, Any]] = []
    warnings: list[str] = []
    for entries in grouped.values():
        entries.sort(
            key=lambda item: (int(item.get("revision") or 0), _clean(item.get("filename"))),
            reverse=True,
        )
        selected.append(entries[0])
        if len(entries) > 1:
            warnings.append(
                f"PO {entries[0]['po_number']}：保留最新修订 {entries[0]['filename']}，忽略 "
                + "、".join(_clean(item.get("filename")) for item in entries[1:])
            )
    for document in ignored:
        warnings.extend(
            f"{document.get('filename')}：{message}"
            for message in document.get("warnings", [])
        )
    rows = [dict(row) for document in selected for row in document.get("rows", [])]
    return rows, warnings


def _grand_total_row(worksheet) -> int:
    candidates = [
        row_no
        for row_no in range(4, (worksheet.max_row or 4) + 1)
        if _label_key(worksheet.cell(row_no, 5).value) in {"合计", "合計", "總計", "总计"}
    ]
    if not candidates:
        raise ValueError("ITEM表未找到末尾合计行")
    return candidates[-1]


def read_schedule(
    source: str | Path | bytes | BinaryIO,
    *,
    filename: str,
) -> dict[str, Any]:
    workbook = new_order_excel.load_complete_workbook_compatible(source, filename=filename)
    try:
        missing = [name for name in (ITEM_SHEET, REVIEW_SHEET, ORDER_SHEET) if name not in workbook.sheetnames]
        if missing:
            raise ValueError("迪士尼排期缺少工作表：" + "、".join(missing))
        item = workbook[ITEM_SHEET]
        required_headers = {
            2: "客出单日期", 3: "PO号", 4: "客名", 5: "產品編號",
            9: "PO数量", 10: "外箱装箱数", 24: "订单单价USD",
        }
        invalid = [
            title for column, title in required_headers.items()
            if _label_key(item.cell(3, column).value) != _label_key(title)
        ]
        if invalid:
            raise ValueError("ITEM表第3行字段不匹配：" + "、".join(invalid))
        grand_total = _grand_total_row(item)
        records: list[dict[str, Any]] = []
        names: dict[str, dict[str, set[str]]] = defaultdict(lambda: {"en": set(), "zh": set()})
        for row_no in range(4, grand_total):
            po_number = _clean(item.cell(row_no, 3).value)
            item_no = _clean(item.cell(row_no, 5).value)
            if not po_number or not item_no:
                continue
            record = {
                "po_number": po_number,
                "item": item_no,
                "quantity": item.cell(row_no, 9).value,
                "customer": _clean(item.cell(row_no, 4).value),
                "row": row_no,
            }
            records.append(record)
            english = _clean(item.cell(row_no, 6).value)
            chinese = _clean(item.cell(row_no, 7).value)
            if english:
                names[_key(item_no)]["en"].add(english)
            if chinese:
                names[_key(item_no)]["zh"].add(chinese)
        item_master = {
            item_no: {
                "product_name_en": next(iter(values["en"])) if len(values["en"]) == 1 else "",
                "product_name_zh": next(iter(values["zh"])) if len(values["zh"]) == 1 else "",
                "name_conflict": len(values["zh"]) > 1,
            }
            for item_no, values in names.items()
        }
        return {
            "records": records,
            "item_master": item_master,
            "sheet": ITEM_SHEET,
            "grand_total_row": grand_total,
        }
    finally:
        workbook.close()


def add_schedule_fields(record: dict[str, Any], schedule: Mapping[str, Any]) -> dict[str, Any]:
    master = schedule.get("item_master", {}).get(_key(record.get("item")), {})
    if not _clean(record.get("description")):
        record["description"] = master.get("product_name_en") or ""
    record["product_name_zh"] = master.get("product_name_zh") or record.get("product_name_zh") or ""
    quantity = _number(record.get("quantity"))
    case_pack = _number(record.get("case_pack"))
    price = _number(record.get("unit_price_usd"))
    record["cartons"] = (
        float(Decimal(str(quantity)) / Decimal(str(case_pack)))
        if quantity not in (None, "") and case_pack not in (None, "", 0)
        else None
    )
    record["inspection_date"] = _minus_days(record.get("ship_date"), 5)
    if price not in (None, ""):
        usd = Decimal(str(price))
        hkd = usd * EXCHANGE_RATE
        record["unit_price_hkd"] = float(hkd)
        if quantity not in (None, ""):
            record["amount_usd"] = float(usd * Decimal(str(quantity)))
            record["amount_hkd"] = float(hkd * Decimal(str(quantity)))
    flags: list[dict[str, Any]] = []
    required = (
        ("po_number", "P/O#"), ("customer", "客名"), ("item", "产品编号"),
        ("description", "英文产品名称"), ("product_name_zh", "中文产品名称"),
        ("quantity", "PO数量"), ("case_pack", "外箱装箱数"),
        ("ship_date", "走货期"), ("unit_price_usd", "订单单价 USD"),
    )
    for field, label in required:
        if record.get(field) in (None, "", 0):
            flags.append({
                "level": "high",
                "code": f"missing_{field}",
                "field": field,
                "text": f"缺少{label}",
            })
    if master.get("name_conflict"):
        flags.append({
            "level": "high",
            "code": "conflicting_product_name_zh",
            "field": "product_name_zh",
            "text": "当前排期同一货号存在多个中文品名，请人工确认",
        })
    if not _clean(record.get("date_code")):
        flags.append({
            "level": "high",
            "code": "missing_date_code",
            "field": "date_code",
            "text": "PO不含完期，日期码需人工填入或确认暂缺",
        })
    if record.get("factory_unit_price_hkd") in (None, ""):
        flags.append({
            "level": "high",
            "code": "missing_factory_unit_price_hkd",
            "field": "factory_unit_price_hkd",
            "text": "出厂价需人工填入或确认暂缺",
        })
    record["flags"] = flags
    record["risk_level"] = "high" if any(flag["level"] == "high" for flag in flags) else "low"
    return record


def _copy_row_style(worksheet, source_row: int, target_row: int, max_col: int) -> None:
    source_dimension = worksheet.row_dimensions[source_row]
    worksheet.row_dimensions[target_row].height = source_dimension.height
    for column in range(1, max_col + 1):
        source = worksheet.cell(source_row, column)
        target = worksheet.cell(target_row, column)
        target.value = None
        target._style = copy(source._style)
        target.comment = None
        target.hyperlink = None


def _insert_rows_preserving(workbook, worksheet, row_no: int, amount: int) -> None:
    if amount <= 0:
        return
    formulas = [
        (sheet, cell.row, cell.column, cell.value)
        for sheet in workbook.worksheets
        for cell in sheet._cells.values()
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    moved_merges = [
        CellRange(str(value))
        for value in worksheet.merged_cells.ranges
        if value.max_row >= row_no
    ]
    for merged in moved_merges:
        worksheet.unmerge_cells(str(merged))
    worksheet.insert_rows(row_no, amount)
    new_order_excel._shift_target_sheet_structures(
        worksheet,
        row_no,
        amount,
        moved_merges,
    )
    for sheet, old_row, column, formula in formulas:
        destination_row = old_row + amount if sheet is worksheet and old_row >= row_no else old_row
        sheet.cell(destination_row, column).value = new_order_excel._rewrite_formula_for_insert(
            formula,
            formula_sheet=sheet.title,
            target_sheet=worksheet.title,
            formula_row=old_row,
            insert_row=row_no,
            amount=amount,
        )


def _item_group(worksheet, item_no: Any, grand_total: int) -> tuple[list[int], int | None]:
    wanted = _key(item_no)
    rows = [
        row_no
        for row_no in range(4, grand_total)
        if _key(worksheet.cell(row_no, 5).value) == wanted
        and _clean(worksheet.cell(row_no, 3).value)
    ]
    if not rows:
        return [], None
    next_row = rows[-1] + 1
    subtotal = (
        next_row
        if next_row < grand_total
        and _label_key(worksheet.cell(next_row, 7).value) in {"合计", "合計", "總計", "总计"}
        else None
    )
    return rows, subtotal


def _detail_formula_templates(
    worksheet,
    preferred_rows: Sequence[int],
) -> dict[int, tuple[str, str]]:
    grand_total = _grand_total_row(worksheet)
    detail_rows = [
        row_no
        for row_no in range(4, grand_total)
        if _clean(worksheet.cell(row_no, 3).value)
        and _clean(worksheet.cell(row_no, 5).value)
    ]
    candidates = list(dict.fromkeys([*preferred_rows, *reversed(detail_rows)]))
    templates: dict[int, tuple[str, str]] = {}
    for column in (25, 26, 27, 28, 29):
        for row_no in candidates:
            cell = worksheet.cell(row_no, column)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                templates[column] = (cell.value, cell.coordinate)
                break
    return templates


def _write_item_formulas(
    worksheet,
    row_no: int,
    templates: Mapping[int, tuple[str, str]],
) -> None:
    fallbacks = {
        25: f"=X{row_no}*7.75",
        27: f"=I{row_no}*X{row_no}",
        28: f"=I{row_no}*Y{row_no}",
        29: f'=IF(Z{row_no}="","",Z{row_no}*I{row_no})',
    }
    for column in (25, 26, 27, 28, 29):
        template = templates.get(column)
        if template:
            formula, origin = template
            try:
                value = Translator(
                    formula,
                    origin=origin,
                ).translate_formula(worksheet.cell(row_no, column).coordinate)
            except (TypeError, ValueError):
                value = formula
            worksheet.cell(row_no, column).value = value
        elif column in fallbacks:
            worksheet.cell(row_no, column).value = fallbacks[column]


def _write_item_detail(
    worksheet,
    row_no: int,
    record: Mapping[str, Any],
    formula_templates: Mapping[int, tuple[str, str]],
) -> None:
    values = {
        2: record.get("received_date"),
        3: record.get("po_number"),
        4: record.get("customer"),
        5: record.get("item"),
        6: record.get("description"),
        7: record.get("product_name_zh"),
        8: record.get("country"),
        9: record.get("quantity"),
        10: record.get("case_pack"),
        13: record.get("date_code"),
        14: record.get("inspection_date"),
        15: record.get("ship_date"),
        24: record.get("unit_price_usd"),
    }
    for column, value in values.items():
        worksheet.cell(row_no, column).value = new_order_excel._safe_record_value(
            {
                2: "order_date", 14: "inspection_date", 15: "ship_date", 13: "date_code",
            }.get(column, "value"),
            value,
        )
    _write_item_formulas(worksheet, row_no, formula_templates)
    if record.get("factory_unit_price_hkd") not in (None, ""):
        worksheet.cell(row_no, 26).value = new_order_excel._safe_record_value(
            "factory_unit_price_hkd",
            record.get("factory_unit_price_hkd"),
        )


def _write_item_subtotal(worksheet, row_no: int, start_row: int, end_row: int) -> None:
    worksheet.cell(row_no, 7).value = "合计"
    worksheet.cell(row_no, 9).value = f"=SUM(I{start_row}:I{end_row})"


def _append_item_rows(workbook, records: list[dict[str, Any]]) -> tuple[list[int], list[int]]:
    worksheet = workbook[ITEM_SHEET]
    linked_rows: list[int] = []
    detail_rows: list[int] = []
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        grouped[_key(record.get("item"))].append(record)

    existing_groups: list[tuple[int, str, list[dict[str, Any]]]] = []
    new_groups: list[tuple[str, list[dict[str, Any]]]] = []
    grand_total = _grand_total_row(worksheet)
    for item_key, rows in grouped.items():
        source_rows, subtotal = _item_group(worksheet, rows[0].get("item"), grand_total)
        if source_rows:
            existing_groups.append((source_rows[0], item_key, rows))
        else:
            new_groups.append((item_key, rows))

    for _first_row, _item_key_value, group_records in sorted(existing_groups):
        grand_total = _grand_total_row(worksheet)
        source_rows, subtotal = _item_group(worksheet, group_records[0].get("item"), grand_total)
        if not source_rows:
            raise ValueError(f"{group_records[0].get('item')}：未找到既有货号")
        group_records.sort(key=lambda item: (_clean(item.get("inspection_date")), _clean(item.get("po_number"))))
        for record in group_records:
            grand_total = _grand_total_row(worksheet)
            source_rows, subtotal = _item_group(
                worksheet,
                record.get("item"),
                grand_total,
            )
            if not source_rows:
                raise ValueError(f"{record.get('item')}：未找到既有货号")
            po_key = _key(record.get("po_number"))
            same_po_rows = [
                row_no
                for row_no in source_rows
                if _key(worksheet.cell(row_no, 3).value) == po_key
            ]
            if same_po_rows:
                insert_row = same_po_rows[-1] + 1
            else:
                new_date = _clean(record.get("inspection_date"))
                insert_row = next(
                    (
                        row_no
                        for row_no in source_rows
                        if new_date
                        and _schedule_date(worksheet.cell(row_no, 14).value)
                        and _schedule_date(worksheet.cell(row_no, 14).value) > new_date
                    ),
                    subtotal if subtotal is not None else source_rows[-1] + 1,
                )
            style_row = next(
                (row_no for row_no in reversed(source_rows) if row_no < insert_row),
                source_rows[0],
            )
            formula_templates = _detail_formula_templates(
                worksheet,
                [style_row, *reversed(source_rows)],
            )
            _insert_rows_preserving(workbook, worksheet, insert_row, 1)
            shifted_style_row = style_row + 1 if style_row >= insert_row else style_row
            _copy_row_style(worksheet, shifted_style_row, insert_row, 29)
            _write_item_detail(
                worksheet,
                insert_row,
                record,
                formula_templates,
            )
            row_no = insert_row
            linked_rows.append(row_no)
            detail_rows.append(row_no)

    if new_groups:
        grand_total = _grand_total_row(worksheet)
        last_used = max(
            (
                row_no
                for row_no in range(4, grand_total)
                if any(worksheet.cell(row_no, column).value not in (None, "") for column in range(2, 30))
            ),
            default=3,
        )
        append_row = last_used + 1
        layout_size = sum(len(rows) + 1 for _key_value, rows in new_groups)
        available = max(0, grand_total - append_row)
        shortfall = max(0, layout_size - available)
        if shortfall:
            _insert_rows_preserving(workbook, worksheet, grand_total, shortfall)
            grand_total += shortfall
        detail_style_row = next(
            (
                row_no
                for row_no in range(append_row - 1, 3, -1)
                if _clean(worksheet.cell(row_no, 3).value) and _clean(worksheet.cell(row_no, 5).value)
            ),
            4,
        )
        subtotal_style_row = next(
            (
                row_no
                for row_no in range(append_row - 1, 3, -1)
                if _label_key(worksheet.cell(row_no, 7).value) in {"合计", "合計", "總計", "总计"}
            ),
            detail_style_row,
        )
        detail_formula_templates = _detail_formula_templates(
            worksheet,
            [detail_style_row],
        )
        row_no = append_row
        for _item_key_value, group_records in new_groups:
            group_records.sort(key=lambda item: (_clean(item.get("inspection_date")), _clean(item.get("po_number"))))
            start_row = row_no
            for record in group_records:
                _copy_row_style(worksheet, detail_style_row, row_no, 29)
                _write_item_detail(
                    worksheet,
                    row_no,
                    record,
                    detail_formula_templates,
                )
                linked_rows.append(row_no)
                detail_rows.append(row_no)
                row_no += 1
            _copy_row_style(worksheet, subtotal_style_row, row_no, 29)
            _write_item_subtotal(worksheet, row_no, start_row, row_no - 1)
            linked_rows.append(row_no)
            row_no += 1

    grand_total = _grand_total_row(worksheet)
    worksheet["I2"] = f"=SUM(I4:I{grand_total - 1})"
    return detail_rows, linked_rows


def _linked_item_formula_rows(worksheet) -> list[tuple[int, int]]:
    linked: list[tuple[int, int]] = []
    pattern = re.compile(rf"(?:'{re.escape(ITEM_SHEET)}'|{re.escape(ITEM_SHEET)})!\$?C\$?(\d+)", re.I)
    for row_no in range(2, (worksheet.max_row or 2) + 1):
        formula = worksheet.cell(row_no, 3).value
        if not isinstance(formula, str) or not formula.startswith("="):
            continue
        match = pattern.search(formula)
        if match:
            linked.append((row_no, int(match.group(1))))
    return linked


def _matching_link_style_row(
    workbook,
    worksheet,
    links: Sequence[tuple[int, int]],
    item_row: int,
) -> int:
    item_sheet = workbook[ITEM_SHEET]
    target_is_subtotal = _label_key(item_sheet.cell(item_row, 7).value) in {
        "合计", "合計", "總計", "总计",
    }
    same_kind = [
        (sheet_row, linked_item_row)
        for sheet_row, linked_item_row in links
        if (
            _label_key(item_sheet.cell(linked_item_row, 7).value)
            in {"合计", "合計", "總計", "总计"}
        ) == target_is_subtotal
    ]
    candidates = same_kind or list(links)
    if not candidates:
        return 2
    return min(
        candidates,
        key=lambda value: (abs(value[1] - item_row), value[1] > item_row, value[0]),
    )[0]


def _append_review_rows(workbook, item_rows: Sequence[int]) -> None:
    if not item_rows:
        return
    worksheet = workbook[REVIEW_SHEET]
    mapping = {2: 2, 3: 3, 4: 4, 5: 5, 6: 9, 7: 10, 8: 15, 10: 25, 11: 29}
    for item_row in sorted(item_rows):
        links = _linked_item_formula_rows(worksheet)
        next_link = next(
            ((sheet_row, linked_item_row) for sheet_row, linked_item_row in links if linked_item_row > item_row),
            None,
        )
        row_no = next_link[0] if next_link else (links[-1][0] + 1 if links else 2)
        style_row = _matching_link_style_row(
            workbook,
            worksheet,
            links,
            item_row,
        )
        if next_link or any(
            worksheet.cell(row_no, column).value not in (None, "")
            for column in range(1, 12)
        ):
            _insert_rows_preserving(workbook, worksheet, row_no, 1)
            if style_row >= row_no:
                style_row += 1
        _copy_row_style(worksheet, style_row, row_no, 11)
        for target_column, item_column in mapping.items():
            item_letter = workbook[ITEM_SHEET].cell(item_row, item_column).column_letter
            worksheet.cell(row_no, target_column).value = f"={ITEM_SHEET}!{item_letter}{item_row}"


def _append_order_rows(workbook, item_rows: Sequence[int]) -> None:
    if not item_rows:
        return
    worksheet = workbook[ORDER_SHEET]
    mapping = {
        1: 2, 3: 3, 5: 3, 6: 4, 7: 8, 8: 5, 10: 7, 11: 9,
        38: 14, 39: 15, 40: 25, 43: 10, 44: 13,
    }
    for item_row in sorted(item_rows):
        links = _linked_item_formula_rows(worksheet)
        summary_row = next(
            (
                row_no
                for row_no in range(2, (worksheet.max_row or 2) + 1)
                if "合计" in _clean(worksheet.cell(row_no, 10).value)
            ),
            (worksheet.max_row or 1) + 1,
        )
        next_link = next(
            ((sheet_row, linked_item_row) for sheet_row, linked_item_row in links if linked_item_row > item_row),
            None,
        )
        row_no = next_link[0] if next_link else summary_row
        style_row = _matching_link_style_row(
            workbook,
            worksheet,
            links,
            item_row,
        )
        _insert_rows_preserving(workbook, worksheet, row_no, 1)
        if style_row >= row_no:
            style_row += 1
        _copy_row_style(worksheet, style_row, row_no, 47)
        for target_column, item_column in mapping.items():
            item_letter = workbook[ITEM_SHEET].cell(item_row, item_column).column_letter
            worksheet.cell(row_no, target_column).value = f"={ITEM_SHEET}!{item_letter}{item_row}"
        worksheet.cell(row_no, 41).value = f"=AN{row_no}*K{row_no}"


def create_export(
    records: Iterable[Mapping[str, Any]],
    output_path: str | Path | BinaryIO,
    schedule_source: str | Path | bytes | BinaryIO,
    *,
    template_filename: str,
) -> dict[str, Any]:
    workbook = new_order_excel.load_complete_workbook_compatible(
        schedule_source,
        filename=template_filename,
    )
    rows = [dict(record) for record in records]
    try:
        missing = [name for name in (ITEM_SHEET, REVIEW_SHEET, ORDER_SHEET) if name not in workbook.sheetnames]
        if missing:
            raise ValueError("迪士尼排期缺少工作表：" + "、".join(missing))
        detail_rows, linked_rows = _append_item_rows(workbook, rows)
        _append_review_rows(workbook, linked_rows)
        _append_order_rows(workbook, linked_rows)
        calculation = getattr(workbook, "calculation", None)
        if calculation is not None:
            calculation.fullCalcOnLoad = True
            calculation.forceFullCalc = True
            calculation.calcMode = "auto"
        if hasattr(output_path, "write"):
            workbook.save(output_path)
            reference = "<memory>"
        else:
            target = Path(output_path)
            target.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(target)
            reference = str(target)
        return {
            "path": reference,
            "detail_rows": detail_rows,
            "linked_rows": linked_rows,
            "mode": "full_workbook_disney_three_sheet_append",
        }
    finally:
        workbook.close()
