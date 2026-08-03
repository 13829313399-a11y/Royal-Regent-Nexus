from __future__ import annotations

import io
import math
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import pdfplumber
import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from .new_order_excel import create_new_order_workbook


FIELDS = {
    "证书": "certificate", "客出单日期": "order_date", "预备单号（OQF NO）": "oqf_no",
    "预备单号(OQF NO)": "oqf_no", "O/C NO": "oc_no", "PO.NO": "po_no",
    "接单日期": "order_date", "客单号": "po_no", "合同": "oc_no", "PO/NO": "po_no",
    "客名/國家": "customer", "客名/国家": "customer", "產品編號": "item_no", "产品编号": "item_no",
    "客名": "customer", "小货号": "item_no", "大货号": "customer_item",
    "產品名称": "product_name", "產品名稱": "product_name", "产品名称": "product_name",
    "数量": "quantity", "數量": "quantity",
    "货名": "product_name", "订单数量": "quantity",
    "装箱": "pack_qty", "箱数": "cartons", "行Q": "factory_inspection", "客Q": "customer_inspection",
    "行Q期": "factory_inspection", "客验期": "customer_inspection",
    "Q货情况": "inspection_status", "客要求走货期": "ship_date", "包装要求": "packaging",
    "订单截数期": "ship_date", "出货期": "ship_date",
    "包装": "packaging", "国家标准": "standard", "备注": "notes", "備註": "notes",
    "生产车间": "workshop", "上系统": "system_status", "单价": "unit_price", "金额HKD": "amount",
    "单价HK$": "unit_price", "金额HK$": "amount",
    "金额": "amount", "客货号": "customer_item", "G.R.N. NO:": "grn_no", "C.R. NO:": "cr_no",
    "外尺碼": "carton_size", "CU.FT": "cuft", "G.W.": "gross_weight", "N.W.": "net_weight",
}
CURRENT_SHEETS = ("手掌", "吊鬼", "衣服", "面具")
SEASONS_SHEETS = ("正单评审表",)
SCHEDULE_FAMILY_LABELS = {
    "seasons": "SEASONS（施信）QF/正式PO排期",
    "internal": "施信综合内部排期",
}
EXPORT_FIELDS = [
    ("order_date", "客出单日期"), ("document_type", "单据状态"), ("oqf_no", "预备单号（OQF NO）"),
    ("oc_no", "O/C NO"), ("po_no", "PO.NO"), ("customer", "客名/国家"), ("item_no", "产品编号"),
    ("product_name", "产品名称"), ("quantity", "数量"), ("remaining_qf", "QF剩余数量"),
    ("pack_qty", "装箱"), ("cartons", "箱数"), ("unit_price", "单价"), ("amount", "金额HKD"),
    ("ship_date", "客要求走货期"), ("packaging", "包装要求"), ("standard", "国家标准"),
    ("inspection_status", "Q货情况"), ("notes", "备注"), ("risk_level", "风险"),
]


def schedule_family_from_sheets(sheet_names: list[str]) -> str:
    """Keep the two schedules sent by merchandising in separate data lanes."""
    names = set(sheet_names)
    if {"手掌", "吊鬼", "衣服", "面具"} & names:
        return "internal"
    if "正单评审表" in names or any(name.endswith("接单表") for name in names):
        return "seasons"
    return "unknown"


def detect_schedule_family(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".xls":
        names = xlrd.open_workbook(file_contents=data, on_demand=True).sheet_names()
    else:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            names = list(book.sheetnames)
        finally:
            book.close()
    return schedule_family_from_sheets(names)


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value)).strip()


def number(value: Any) -> float | int:
    if value in (None, ""):
        return 0
    try:
        result = float(str(value).replace(",", ""))
        return int(result) if result.is_integer() else round(result, 4)
    except (ValueError, TypeError):
        return 0


def excel_date(value: Any, datemode: int = 0) -> str:
    if not value:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)) and 20000 < value < 80000:
        return xlrd.xldate_as_datetime(value, datemode).strftime("%Y-%m-%d")
    text = clean(value).replace("/", "-")
    match = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
    return f"{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}" if match else text


def add_state(record: dict[str, Any]) -> dict[str, Any]:
    record["quantity"] = number(record.get("quantity"))
    record["unit_price"] = number(record.get("unit_price"))
    record["amount"] = number(record.get("amount")) or round(record["quantity"] * record["unit_price"], 2)
    record["document_type"] = record.get("document_type") or ("正式PO" if record.get("po_no") else "QF预备单")
    risks = []
    # 自洽校验：数量×单价 ≠ 单据金额时标红（不静默采信任何一边）
    if record["quantity"] > 0 and record["unit_price"] > 0 and record["amount"] > 0:
        expect = round(record["quantity"] * record["unit_price"], 2)
        if abs(expect - record["amount"]) > 0.05:
            risks.append("金额与数量×单价不符")
    if record["document_type"] == "QF预备单" and record["quantity"] > 0:
        risks.append("待正式PO")
    if record.get("po_no") and not record.get("oc_no"):
        risks.append("缺少O/C NO")
    if not record.get("ship_date"):
        risks.append("缺走货期")
    if record.get("quantity", 0) < 0:
        risks.append("QF超扣")
    record["risk_text"] = "、".join(risks)
    record["risk_level"] = "high" if any(x in risks for x in ("QF超扣", "缺走货期", "金额与数量×单价不符")) else "medium" if risks else "normal"
    return record


def _parse_xls(data: bytes) -> tuple[list[dict[str, Any]], str]:
    book = xlrd.open_workbook(file_contents=data)
    records = []
    parsed_sheets = []
    preferred = [name for name in CURRENT_SHEETS if name in book.sheet_names()]
    candidates = preferred or (["正单评审表"] if "正单评审表" in book.sheet_names() else book.sheet_names())
    for sheet_name in candidates:
        sheet = book.sheet_by_name(sheet_name)
        header_row = next((r for r in range(min(12, sheet.nrows)) if sum(clean(sheet.cell_value(r, c)) in FIELDS for c in range(sheet.ncols)) >= 5), None)
        if header_row is None:
            continue
        mapping = {c: FIELDS[clean(sheet.cell_value(header_row, c))] for c in range(sheet.ncols) if clean(sheet.cell_value(header_row, c)) in FIELDS}
        parsed_sheets.append(sheet.name)
        for r in range(header_row + 1, sheet.nrows):
            row = {field: sheet.cell_value(r, c) for c, field in mapping.items()}
            if not any(clean(row.get(k)) for k in ("po_no", "item_no", "customer")) or clean(row.get("oqf_no")).startswith("↑"):
                continue
            if clean(row.get("product_name")) in {"总合计", "余数"}:
                continue
            row["order_date"] = excel_date(row.get("order_date"), book.datemode)
            row["ship_date"] = excel_date(row.get("ship_date"), book.datemode)
            row = {k: clean(v) if k not in {"quantity", "cartons", "unit_price", "amount", "cuft", "gross_weight", "net_weight"} else number(v) for k, v in row.items()}
            row["schedule_category"] = sheet.name
            records.append(add_state(row))
    if not parsed_sheets:
        raise ValueError("未识别到施信客排期表表头")
    return records, "、".join(parsed_sheets)


def _parse_xlsx(data: bytes) -> tuple[list[dict[str, Any]], str]:
    book = load_workbook(io.BytesIO(data), data_only=True)
    records = []
    parsed_sheets = []
    preferred = [name for name in CURRENT_SHEETS if name in book.sheetnames]
    candidates = preferred or (["正单评审表"] if "正单评审表" in book.sheetnames else book.sheetnames)
    for sheet_name in candidates:
        sheet = book[sheet_name]
        header_row = None
        mapping = {}
        for r in range(1, min(sheet.max_row, 12) + 1):
            current = {c: FIELDS[clean(sheet.cell(r, c).value)] for c in range(1, sheet.max_column + 1) if clean(sheet.cell(r, c).value) in FIELDS}
            if len(current) >= 5:
                header_row, mapping = r, current
                break
        if header_row is None:
            continue
        parsed_sheets.append(sheet.title)
        for r in range(header_row + 1, sheet.max_row + 1):
            row = {field: sheet.cell(r, c).value for c, field in mapping.items()}
            if not any(clean(row.get(k)) for k in ("po_no", "item_no", "customer")) or clean(row.get("oqf_no")).startswith("↑"):
                continue
            if clean(row.get("product_name")) in {"总合计", "余数"}:
                continue
            row["order_date"] = excel_date(row.get("order_date"))
            row["ship_date"] = excel_date(row.get("ship_date"))
            row = {k: clean(v) if k not in {"quantity", "cartons", "unit_price", "amount", "cuft", "gross_weight", "net_weight"} else number(v) for k, v in row.items()}
            row["schedule_category"] = sheet.title
            records.append(add_state(row))
    if not parsed_sheets:
        raise ValueError("未识别到施信客排期表表头")
    return records, "、".join(parsed_sheets)


def parse_schedule(data: bytes, filename: str) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    records, sheet = _parse_xls(data) if suffix == ".xls" else _parse_xlsx(data)
    family = detect_schedule_family(data, filename)
    return {
        "filename": filename,
        "sheet": sheet,
        "family": family,
        "family_label": SCHEDULE_FAMILY_LABELS.get(family, "未识别排期类别"),
        "records": records,
        "count": len(records),
    }


def _match(pattern: str, text: str, flags: int = re.I) -> str:
    found = re.search(pattern, text, flags)
    return clean(found.group(1)) if found else ""


def _pdf_text(data: bytes) -> tuple[str, int]:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = [page.extract_text() or "" for page in pdf.pages]
    return "\n".join(pages), len(pages)


def _cell_text(value: Any) -> str:
    if value in (None, ""):
        return ""
    text = re.sub(r"[ \t]+", " ", str(value).replace("\r\n", "\n")).strip()
    # WPS occasionally splits the decimal point into a separately positioned glyph.
    return re.sub(r"(?<=\d)\s+\.(?=\d)", ".", text)


def _meaningful_cell(value: str) -> bool:
    # WPS PDF-to-Excel conversions add many layout cells containing numeric 0.
    return bool(value) and not re.fullmatch(r"0(?:\.0+)?", value)


def _workbook_content(data: bytes, suffix: str) -> tuple[str, int, list[str]]:
    """Read WPS conversions by row so split cells still form one logical PO line."""
    rows: list[str] = []
    if suffix == ".xls":
        book = xlrd.open_workbook(file_contents=data)
        for sheet in book.sheets():
            for row_index in range(sheet.nrows):
                values = [
                    _cell_text(sheet.cell_value(row_index, column_index))
                    for column_index in range(sheet.ncols)
                ]
                values = [value for value in values if _meaningful_cell(value)]
                if values:
                    rows.append(" ".join(values))
        return "\n".join(rows), len(book.sheets()), rows

    book = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        for sheet in book.worksheets:
            for row in sheet.iter_rows(values_only=True):
                values = [_cell_text(value) for value in row]
                values = [value for value in values if _meaningful_cell(value)]
                if values:
                    rows.append(" ".join(values))
        return "\n".join(rows), len(book.worksheets), rows
    finally:
        book.close()


def _filename_customer(filename: str) -> str:
    stem = Path(filename).stem
    stem = re.sub(r"-signed.*$", "", stem, flags=re.I)
    parts = [part.strip() for part in stem.split(" - ")]
    if len(parts) >= 4:
        return re.sub(r"-R\d+$", "", " - ".join(parts[3:]), flags=re.I).strip()
    return ""


def _compact_document_code(pattern: str, text: str) -> str:
    found = re.search(pattern, text, re.I)
    return re.sub(r"[^A-Z0-9-]", "", found.group(1).upper()) if found else ""


def _document_header(text: str, filename: str, is_qf: bool) -> dict[str, str]:
    po_no = _compact_document_code(r"\b((?:M|Z)\s*P\s*O(?:\s*\d){6,})", filename)
    if not po_no:
        po_no = _compact_document_code(r"\b((?:M|Z)\s*P\s*O(?:\s*\d){6,})", text)
    # 文件名常写成 “QF16098161 - 247 - 客户名”，其中 247 是跟单/地区标记，
    # 不是 QF 号后缀；只有正文 OQF No 明确带出的 “-01” 等后缀才保留。
    oqf_no = _compact_document_code(r"\b(Q\s*F(?:\s*\d){6,})", filename)
    if not oqf_no:
        oqf_no = _compact_document_code(
            r"O\s*Q\s*F\s*(?:No\.?|NO\.?)?\s*[:：]?\s*(Q\s*F(?:\s*\d){6,}(?:\s*-\s*\d+)?)",
            text,
        )
    customer = _match(r"Customer\s+Name\s*:\s*(.+?)(?:\s+Order\s+Date\s*:|\n)", text)
    customer = re.sub(r"^C\d+\s*", "", customer, flags=re.I).strip()
    customer = _filename_customer(filename) or customer
    order_date = excel_date(
        _match(r"(?:Order\s+)?Date\s*:\s*(\d{4}/\d{1,2}/\d{1,2})", text)
    )
    return {
        "document_type": "QF预备单" if is_qf else "正式PO",
        "po_no": "" if is_qf else po_no,
        "oqf_no": oqf_no,
        "customer": customer,
        "order_date": order_date,
    }


_FORMAL_LINE = re.compile(
    r"(?m)^\s*(?P<item>[A-Z][A-Z0-9-]{3,})\s+"
    r"OC\s*#?\s*\.?\s*:?\s*(?P<oc>[A-Z]{2}\d+)\s+"
    r"SHIPDATE\s*:\s*(?P<ship>\d{4}/\d{1,2}/\d{1,2})\s+"
    r"(?P<qty>[\d,]+)\s+(?P<price>[\d,]+(?:\.\d+)?)\s+"
    r"(?:\d+(?:\.\d+)?%?\s+)?(?P<amount>[\d,]+(?:\.\d+)?)",
    re.I,
)
_QF_LINE = re.compile(
    r"(?m)^\s*(?P<item>[A-Z][A-Z0-9-]{3,})\s+"
    r"SHIPDATE\s*:\s*(?P<ship>\d{4}/\d{1,2}/\d{1,2})\s+"
    r"(?P<qty>[\d,]+)(?:\s|$)",
    re.I,
)


def _product_name_from_block(block: str) -> str:
    lines = [re.sub(r"\s+", " ", line).strip() for line in block.splitlines()]
    ignored = (
        "PACKING", "ITEM NO", "OQF", "CUSTOMER", "TOTAL", "SHIPDATE",
        "OC #", "CTN MEAS", "CUBE", "OUTER", "INNER",
    )
    for line in lines[1:12]:
        if not line or line.upper().startswith(ignored):
            continue
        candidate = re.split(r"\s+\d+\.", line, maxsplit=1)[0].strip()
        if re.fullmatch(r"[A-Z0-9][A-Z0-9 '&/+,.()\"-]{4,}", candidate):
            return candidate
    return ""


def _packing_from_block(block: str) -> tuple[str, str]:
    pack_qty = _match(r"Outer\s*(?:Type\s*:\s*\w+\s*)?(?:QTY\s*:|:)\s*([\d,]+)\s*PCS", block)
    if not pack_qty:
        pack_qty = _match(r"Outer\s*:\s*([\d,]+)\s*PCS", block)
    packaging = _match(r"(Packing\s*(?:/ Cube)?\s*:[^\n]+)", block)
    if packaging and not re.search(
        r"(?:\d+\s*PCS|POLYBAG|CARTON|PDQ|DISPLAY|BOX|CARD|袋|箱|卡|彩贴)",
        packaging,
        re.I,
    ):
        packaging = ""
    return pack_qty.replace(",", ""), packaging


def _workbook_line_parts(row: str, is_qf: bool) -> dict[str, str] | None:
    if is_qf:
        match = re.search(
            r"(?s)(?P<item>[A-Z][A-Z0-9-]{3,}).*?"
            r"SHIPDATE\s*:\s*(?P<ship>\d{4}/\d{1,2}/\d{1,2})\s+"
            r"(?P<qty>[\d,]+)",
            row,
            re.I,
        )
        return match.groupdict() if match else None
    head = re.search(
        r"(?s)(?P<item>[A-Z][A-Z0-9-]{3,}).*?"
        r"OC\s*#?\s*\.?\s*:?\s*(?P<oc>[A-Z]{2}\d+).*?"
        r"SHIPDATE\s*:\s*(?P<ship>\d{4}/\d{1,2}/\d{1,2})",
        row,
        re.I,
    )
    if not head:
        return None
    tail = row[head.end():]
    qty = re.search(r"(?<![\d.])([\d,]+)(?![\d.])", tail)
    price_amount = list(
        re.finditer(
            r"([\d,]+\.\d{1,4})\s+\d+(?:\.\d+)?%?\s+([\d,]+(?:\.\d{1,4})?)",
            tail,
        )
    )
    if not qty:
        return None
    if price_amount:
        price, amount = price_amount[-1].group(1), price_amount[-1].group(2)
    else:
        # WPS 转换后的部分 PO 会把折扣 0 隐去，并把数量、单价、金额拆成
        # 独立单元格，例如 “101 101 175.5 17725.5”。用单据内部的
        # 数量×单价=金额自洽关系挑选单价和金额，避免把重复数量误认成价格。
        quantity = number(qty.group(1))
        numeric_tokens = [
            token
            for token in re.findall(r"(?<![A-Z0-9])([\d,]+(?:\.\d+)?)(?![A-Z0-9])", tail, re.I)
            if number(token) > 0
        ]
        arithmetic_matches: list[tuple[float, int, str, str]] = []
        for price_index in range(1, len(numeric_tokens) - 1):
            candidate_price = number(numeric_tokens[price_index])
            if candidate_price <= 0:
                continue
            for amount_index in range(price_index + 1, len(numeric_tokens)):
                candidate_amount = number(numeric_tokens[amount_index])
                error = abs(round(quantity * candidate_price, 2) - candidate_amount)
                if error <= 0.05:
                    arithmetic_matches.append(
                        (
                            error,
                            amount_index - price_index,
                            numeric_tokens[price_index],
                            numeric_tokens[amount_index],
                        )
                    )
        if not arithmetic_matches:
            return None
        _, _, price, amount = min(arithmetic_matches)
    return {
        **head.groupdict(),
        "qty": qty.group(1),
        "price": price,
        "amount": amount,
    }


def _parse_document_lines(
    text: str,
    header: dict[str, str],
    is_qf: bool,
    filename: str,
    workbook_rows: list[str] | None = None,
) -> list[dict[str, Any]]:
    candidates: list[tuple[dict[str, str], str]] = []
    if workbook_rows:
        for row in workbook_rows:
            parts = _workbook_line_parts(row, is_qf)
            if parts:
                candidates.append((parts, row))
    if not candidates:
        pattern = _QF_LINE if is_qf else _FORMAL_LINE
        matches = list(pattern.finditer(text))
        for index, match in enumerate(matches):
            block_end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
            candidates.append((match.groupdict(), text[match.start():block_end]))
    records: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str, str]] = set()
    for index, (parts, block) in enumerate(candidates):
        dedupe = (
            parts.get("item", "").upper(),
            parts.get("oc", "").upper(),
            parts.get("ship", ""),
            parts.get("qty", ""),
        )
        if dedupe in seen:
            continue
        seen.add(dedupe)
        pack_qty, packaging = _packing_from_block(block)
        record: dict[str, Any] = {
            **header,
            "source_file": filename,
            "source_row": index + 1,
            "item_no": parts["item"].upper(),
            "source_item_no": parts["item"].upper(),
            "oc_no": "" if is_qf else parts["oc"].upper(),
            "ship_date": excel_date(parts["ship"]),
            "quantity": number(parts["qty"]),
            "unit_price": 0 if is_qf else number(parts["price"]),
            "amount": 0 if is_qf else number(parts["amount"]),
            "english_name": _product_name_from_block(block),
            "product_name": _product_name_from_block(block),
            "pack_qty": number(pack_qty) if pack_qty else "",
            "cartons": 0,
            "packaging": packaging,
            "standard": "",
            "notes": "",
            "schedule_family": "seasons",
        }
        if record["quantity"] and record["pack_qty"]:
            record["cartons"] = math.ceil(record["quantity"] / number(record["pack_qty"]))
        records.append(record)
    return records


def _item_key(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", clean(value).upper())


def _qf_base(value: Any) -> str:
    return re.sub(r"-\d+$", "", clean(value).upper())


def _schedule_match(record: dict[str, Any], schedule_records: list[dict[str, Any]]) -> dict[str, Any] | None:
    po_no = clean(record.get("po_no")).upper()
    oqf_no = _qf_base(record.get("oqf_no"))
    source_item = _item_key(record.get("item_no"))
    candidates = []
    for schedule in schedule_records:
        if po_no:
            if clean(schedule.get("po_no")).upper() != po_no:
                continue
        elif oqf_no:
            if _qf_base(schedule.get("oqf_no")) != oqf_no:
                continue
        else:
            continue
        target_item = _item_key(schedule.get("item_no"))
        if source_item == target_item:
            candidates.append((2, schedule))
        elif source_item and target_item and (
            target_item.startswith(source_item) or source_item.startswith(target_item)
        ):
            candidates.append((1, schedule))
    if not candidates:
        return None
    best_score = max(score for score, _ in candidates)
    best = [schedule for score, schedule in candidates if score == best_score]
    if oqf_no:
        # 优先继承尚未转正式 PO 的 QF 预留行；如果排期里只剩正式 PO 行，
        # 仍可用同 QF/货号继承客名、品名等静态字段。
        preorder = [row for row in best if not clean(row.get("po_no"))]
        if preorder:
            best = preorder
    if po_no:
        unique = {
            (
                clean(row.get("po_no")).upper(),
                _item_key(row.get("item_no")),
                excel_date(row.get("ship_date")),
            ): row
            for row in best
        }
    else:
        # 一个基础 QF 号在排期里可能被拆成 -01/-04 等多条预留行。
        # 只要这些候选指向同一货号、同一产品与同一客户，就可安全继承；
        # 不能把排期内部的拆分后缀误判为“找不到 QF”。
        unique = {
            (
                _item_key(row.get("item_no")),
                clean(row.get("product_name")).upper(),
                clean(row.get("customer")).upper(),
            ): row
            for row in best
        }
    return next(iter(unique.values())) if len(unique) == 1 else None


def _enrich_order_records(
    records: list[dict[str, Any]],
    schedule_records: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[str]]:
    warnings: list[str] = []
    fields_to_inherit = (
        "oqf_no", "customer", "product_name", "pack_qty", "cartons", "packaging",
        "standard", "notes", "workshop", "schedule_category",
    )
    for record in records:
        matched = _schedule_match(record, schedule_records)
        record["matched_schedule"] = bool(matched)
        if not matched:
            warnings.append(
                f"{record.get('po_no') or record.get('oqf_no') or '单据'} / "
                f"{record.get('item_no') or '未知货号'}：当前SEASONS排期未找到同单同货号，按新单处理"
            )
            continue
        source_item = clean(record.get("item_no"))
        canonical_item = clean(matched.get("item_no"))
        if source_item and canonical_item and _item_key(source_item) != _item_key(canonical_item):
            warnings.append(f"{source_item} 按当前排期对应为 {canonical_item}")
        record["item_no"] = canonical_item or source_item
        for field in fields_to_inherit:
            value = matched.get(field)
            if field == "product_name":
                if value:
                    record[field] = value
                continue
            if not record.get(field) and value not in (None, ""):
                record[field] = value
        schedule_qty = number(matched.get("quantity"))
        if record.get("document_type") == "正式PO" and schedule_qty and record.get("quantity"):
            if abs(number(record["quantity"]) - schedule_qty) > 0.0001:
                warnings.append(
                    f"{record.get('po_no')} / {record.get('item_no')}："
                    f"PO数量 {record.get('quantity')} 与当前排期 {schedule_qty} 不一致，按修改单拦截复核"
                )
                record["schedule_quantity_conflict"] = True
        for field in ("unit_price", "amount"):
            if not number(record.get(field)) and number(matched.get(field)):
                record[field] = number(matched.get(field))
        pack_qty = number(record.get("pack_qty"))
        if record.get("quantity") and pack_qty > 0 and not record.get("cartons"):
            record["cartons"] = math.ceil(number(record["quantity"]) / pack_qty)
    return records, list(dict.fromkeys(warnings))


def _critical_errors(record: dict[str, Any]) -> list[str]:
    required = [
        ("item_no", "产品编号"),
        ("quantity", "数量"),
        ("ship_date", "走货期"),
        ("customer", "客名"),
        ("product_name", "产品名称"),
    ]
    if record.get("document_type") == "正式PO":
        required[:0] = [("po_no", "PO号"), ("oc_no", "O/C NO")]
    else:
        required.insert(0, ("oqf_no", "预备单号"))
    missing = [label for field, label in required if not record.get(field)]
    if record.get("schedule_quantity_conflict"):
        missing.append("数量与当前排期冲突")
    return missing


def parse_order_documents(
    data: bytes,
    filename: str,
    schedule_records: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text, pages = _pdf_text(data)
        workbook_rows = None
    elif suffix in {".xlsx", ".xlsm", ".xls"}:
        text, pages, workbook_rows = _workbook_content(data, suffix)
    else:
        raise ValueError(f"施信 PO 不支持 {suffix or '无扩展名'}")
    is_qf = "ORDER QUICK FORM" in text.upper() or "QUICK FORM" in text.upper() or bool(
        re.search(r"\bQF\d{6,}", filename, re.I)
    )
    kind = "QF预备单" if is_qf else "正式PO"
    header = _document_header(text, filename, is_qf)
    records = _parse_document_lines(text, header, is_qf, filename, workbook_rows)
    warnings: list[str] = []
    if schedule_records:
        records, inherited_warnings = _enrich_order_records(records, schedule_records)
        warnings.extend(inherited_warnings)
    valid: list[dict[str, Any]] = []
    rejected: list[dict[str, Any]] = []
    for record in records:
        errors = _critical_errors(record)
        if errors:
            rejected.append({**record, "errors": errors})
            warnings.append(
                f"{filename} / {record.get('item_no') or '未识别货号'}："
                f"缺少或冲突字段（{'、'.join(errors)}），本行未生成且未入库"
            )
            continue
        valid.append(add_state(record))
    if not records:
        warnings.append(f"{filename}：未识别到施信{kind}产品明细行，本文件未生成且未入库")
    status = "success" if valid and not rejected else "partial" if valid else "failed"
    return {
        "filename": filename,
        "type": kind,
        "pages": pages,
        "status": status,
        "records": valid,
        "rejected_records": rejected,
        "row_count": len(valid),
        "warnings": list(dict.fromkeys(warnings)),
        "error": "" if valid else "未形成可安全生成的新单明细",
    }


def reconcile_documents(documents: list[dict[str, Any]], existing: list[dict[str, Any]]) -> dict[str, Any]:
    def identity(record: dict[str, Any]) -> tuple[str, str, str]:
        item = clean(record.get("item_no")).upper()
        if record.get("po_no"):
            return "PO", clean(record.get("po_no")).upper(), item
        return "QF", clean(record.get("oqf_no")).upper(), item

    # 解析阶段已确认“同单同货号但数量冲突”的行不会进入安全 records，
    # 但仍须在批量汇总中明确列为修改/异常，不能只藏在文件级警告里。
    modification_records: list[dict[str, Any]] = []
    rejected_modification_ids: set[tuple[str, str, str]] = set()
    for doc in documents:
        for record in doc.get("rejected_records", []):
            if not record.get("schedule_quantity_conflict"):
                continue
            rid = identity(record)
            if rid in rejected_modification_ids:
                continue
            rejected_modification_ids.add(rid)
            record["risk_level"] = "high"
            record["risk_text"] = "当前排期已有同单同货号但数量不同，按修改单人工处理"
            modification_records.append(record)

    # 同一批次重复上传只计一次。
    incoming: list[dict[str, Any]] = []
    _seen_ids: dict[tuple[str, str, str], int] = {}
    _dup_warnings: list[str] = []
    for doc in documents:
        for r in doc["records"]:
            rid = identity(r)
            if all(rid[1:]):
                if rid in _seen_ids:
                    incoming[_seen_ids[rid]] = r
                    _dup_warnings.append(
                        f"{doc.get('filename', '')} 与之前上传的文件是同一单据（{rid[1]} / {rid[2]}），数量只计一次"
                    )
                    continue
                _seen_ids[rid] = len(incoming)
            incoming.append(r)

    # 当前排期已经存在的同单同货号不再作为新单导出；若数量不同，
    # 视为修改/补单并按用户业务边界拦截，不能静默覆盖。
    existing_by_id: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for record in existing:
        rid = identity(record)
        if all(rid[1:]):
            existing_by_id[rid].append(record)
    new_records: list[dict[str, Any]] = []
    existing_records: list[dict[str, Any]] = []
    for record in incoming:
        rid = identity(record)
        matches = existing_by_id.get(rid, [])
        if not matches:
            new_records.append(record)
            continue
        schedule_quantities = {number(item.get("quantity")) for item in matches}
        incoming_quantity = number(record.get("quantity"))
        if incoming_quantity in schedule_quantities:
            existing_records.append(record)
            continue
        record["risk_level"] = "high"
        record["risk_text"] = "当前排期已有同单同货号但数量不同，按修改单人工处理"
        modification_records.append(record)

    incoming_ids = {identity(record) for record in new_records}
    base = [record for record in existing if identity(record) not in incoming_ids]
    groups: dict[tuple[str, str], dict[str, Any]] = defaultdict(lambda: {"qf": 0, "po": 0})
    for record in [*base, *new_records]:
        key = (clean(record.get("oqf_no")).upper(), clean(record.get("item_no")).upper())
        if not all(key):
            continue
        qty = number(record.get("quantity"))
        if record.get("po_no") or record.get("document_type") == "正式PO":
            groups[key]["po"] += qty
        else:
            groups[key]["qf"] += qty
    warnings = [warning for doc in documents for warning in doc.get("warnings", [])] + _dup_warnings
    for record in existing_records:
        warnings.append(
            f"{record.get('po_no') or record.get('oqf_no')} / {record.get('item_no')}："
            "当前SEASONS排期已存在相同数量，本次不重复生成、不重复入库"
        )
    for record in modification_records:
        warnings.append(
            f"{record.get('po_no') or record.get('oqf_no')} / {record.get('item_no')}："
            f"当前排期数量与本次 {record.get('quantity')} 不同，疑似修改/补单，已拦截并交人工处理"
        )
    for record in new_records:
        key = (clean(record.get("oqf_no")).upper(), clean(record.get("item_no")).upper())
        state = groups.get(key, {"qf": 0, "po": 0})
        record["qf_quantity"] = state["qf"]
        record["po_total"] = state["po"]
        record["remaining_qf"] = state["qf"] - state["po"]
        if record["remaining_qf"] < 0:
            record["risk_level"], record["risk_text"] = "high", "QF数量不足，正式PO累计超扣"
            warnings.append(f"{record.get('oqf_no')} / {record.get('item_no')} 正式PO累计超出QF {abs(record['remaining_qf'])} 件")
    files = [
        {
            "name": doc.get("filename", ""),
            "status": doc.get("status", "failed"),
            "row_count": doc.get("row_count", len(doc.get("records", []))),
            "error": doc.get("error", ""),
        }
        for doc in documents
    ]
    success_count = sum(1 for item in files if item["status"] == "success")
    partial_count = sum(1 for item in files if item["status"] == "partial")
    failed_count = sum(1 for item in files if item["status"] == "failed")
    message = (
        f"已处理 {len(documents)} 份施信 QF/PO：安全新单 {len(new_records)} 行，"
        f"当前排期已存在 {len(existing_records)} 行，修改/冲突拦截 {len(modification_records)} 行。"
    )
    return {
        "documents": documents,
        "files": files,
        "records": new_records,
        "existing_records": existing_records,
        "modification_records": modification_records,
        "warnings": list(dict.fromkeys(warnings)),
        "summary": summarize(new_records),
        "count": len(new_records),
        "row_count": len(new_records),
        "message": message,
        "meta": {
            "row_count": len(new_records),
            "success_count": success_count,
            "partial_count": partial_count,
            "failed_count": failed_count,
            "existing_count": len(existing_records),
            "modification_count": len(modification_records),
        },
    }


def summarize(records: list[dict[str, Any]]) -> dict[str, Any]:
    formal = [r for r in records if r.get("po_no")]
    qf = [r for r in records if not r.get("po_no")]
    customers = Counter(clean(r.get("customer")) for r in records if clean(r.get("customer")))
    return {"orders": len(records), "qf": len(qf), "formal_po": len(formal),
            "quantity": sum(number(r.get("quantity")) for r in records),
            "amount": round(sum(number(r.get("amount")) for r in records), 2),
            "risk": sum(1 for r in records if r.get("risk_level") in {"high", "medium"}),
            "customers": customers.most_common(6)}


def build_export(
    records: list[dict[str, Any]],
    path: Path,
    template_source: bytes | str | Path | None = None,
    template_filename: str = "schedule.xlsx",
    sheet_names: tuple[str, ...] = CURRENT_SHEETS,
) -> Path:
    if template_source is not None:
        aliases: dict[str, list[str]] = {}
        for label, field in FIELDS.items():
            aliases.setdefault(field, []).append(label)
        for field, label in EXPORT_FIELDS:
            aliases.setdefault(field, []).append(label)
        create_new_order_workbook(
            template_source,
            path,
            records,
            aliases,
            filename=template_filename,
            sheet_names=sheet_names,
            sheet_title="新单",
        )
        return path
    book = Workbook()
    sheet = book.active
    sheet.title = "施信客排期汇总"
    sheet.append([title for _, title in EXPORT_FIELDS])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="183B56")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for record in records:
        sheet.append([record.get(field, "") for field, _ in EXPORT_FIELDS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, len(EXPORT_FIELDS) + 1):
        width = min(42, max(11, max(len(clean(sheet.cell(row, column).value)) for row in range(1, sheet.max_row + 1)) + 2))
        sheet.column_dimensions[get_column_letter(column)].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path
