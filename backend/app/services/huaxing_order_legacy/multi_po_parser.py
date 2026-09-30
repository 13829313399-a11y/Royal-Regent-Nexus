from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path
from typing import Any, BinaryIO

import pdfplumber
from openpyxl import load_workbook

from app.services.customer_order_ocr import read_scanned_order_pdf


CLIENTS = {
    "maxx": "Maxx",
    "shushupapa": "Shushupapa",
    "barter": "Barter",
}


def _clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _number(value: Any) -> float | int | None:
    text = re.sub(r"[^\d.\-]", "", str(value or ""))
    if not text or text in {"-", "."}:
        return None
    try:
        number = float(text)
        return int(number) if number.is_integer() else number
    except ValueError:
        return None


def _iso_date(value: Any) -> str | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = _clean(value).replace("年", "-").replace("月", "-").replace("日", "")
    text = re.sub(r"(?<=\d)(st|nd|rd|th)\b", "", text, flags=re.I)
    text = re.sub(r"\s+,", ",", text)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d %B %Y", "%B %d, %Y", "%b %d, %Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    match = re.search(r"(20\d{2})[./-](\d{1,2})[./-](\d{1,2})", text)
    if match:
        return f"{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    match = re.search(r"(\d{1,2})[./-](\d{1,2})[./-](20\d{2})", text)
    if match:
        return f"{match.group(3)}-{int(match.group(2)):02d}-{int(match.group(1)):02d}"
    match = re.search(
        r"(\d{1,2})\s+([A-Za-z]+),?\s+(20\d{2})",
        text,
        re.I,
    )
    if match:
        return _iso_date(f"{match.group(1)} {match.group(2)} {match.group(3)}")
    return None


def _pdf_bytes(source: str | Path | bytes | BinaryIO) -> bytes:
    if isinstance(source, bytes):
        return source
    if isinstance(source, (str, Path)):
        return Path(source).read_bytes()
    position = source.tell() if hasattr(source, "tell") else None
    payload = source.read()
    if position is not None and hasattr(source, "seek"):
        source.seek(position)
    return payload


def _ocr_contract_page(pytesseract_module, image, language: str) -> str:
    dense = pytesseract_module.image_to_string(
        image,
        lang=language,
        config="--oem 3 --psm 6 -c preserve_interword_spaces=1",
        timeout=30,
    )
    # A sparse second pass over the contract header/customer band is especially
    # useful for distinguishing 8/9 in Walmart PO numbers without doubling the
    # OCR cost for the full page.
    with image.crop((0, int(image.height * 0.16), image.width, int(image.height * 0.48))) as header_band:
        sparse = pytesseract_module.image_to_string(
            header_band,
            lang=language,
            config="--oem 3 --psm 11",
            timeout=15,
        )
    return dense + "\n" + sparse


def read_pdf_pages(
    source: str | Path | bytes | BinaryIO,
    *,
    min_native_chars_per_page: int = 40,
    max_pages: int = 60,
) -> tuple[list[str], bool]:
    """Read native PDF text, falling back to bounded page-preserving OCR."""
    content = _pdf_bytes(source)
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        if len(pdf.pages) > max_pages:
            raise ValueError(f"扫描 PDF 最多支持 {max_pages} 页，当前为 {len(pdf.pages)} 页")
        pages = [(page.extract_text(x_tolerance=2, y_tolerance=3) or "") for page in pdf.pages]
    if not pages:
        raise ValueError("PDF 没有页面")
    native_chars = sum(len(re.sub(r"\s+", "", page)) for page in pages)
    if native_chars >= len(pages) * min_native_chars_per_page:
        return pages, False

    return read_scanned_order_pdf(content, len(pages)), True


def read_pdf(source: str | Path | bytes | BinaryIO) -> tuple[str, int]:
    pages, _ = read_pdf_pages(source)
    return "\n".join(pages), len(pages)


def detect_client(text: str, filename: str = "") -> str:
    haystack = f"{filename}\n{text}".upper()
    if any(token in haystack for token in ("SHUSHUPAPA", "SHU SHU PAPA", "PRICESMART")):
        return "shushupapa"
    if any(token in haystack for token in ("MAXX MARKETING", "PO-HK-", "MAXX")):
        return "maxx"
    if any(token in haystack for token in ("BARTER", "TNT", "WAL-MART", "WALMART")):
        return "barter"
    return ""


def _search(pattern: str, text: str, flags: int = re.I) -> str:
    match = re.search(pattern, text, flags)
    return _clean(match.group(1)) if match else ""


def _base_result(client: str, filename: str, pages: int = 0) -> dict[str, Any]:
    return {
        "client": client,
        "client_name": CLIENTS.get(client, "待识别"),
        "filename": filename,
        "pages": pages,
        "po_number": "",
        "customer_po": "",
        "po_date": None,
        "ship_date": None,
        "inspection_date": None,
        "project": "",
        "project_no": "",
        "customer": CLIENTS.get(client, ""),
        "shipment_note": "",
        "lines": [],
        "warnings": [],
        "confidence": "high",
    }


def _parse_maxx(text: str, filename: str, pages: int) -> dict[str, Any]:
    result = _base_result("maxx", filename, pages)
    result.update({
        "po_number": _search(r"P\.\s*O\.\s*NO\.\s*([^\s(]+)", text),
        "po_date": _iso_date(_search(r"P\.\s*O\.\s*DATE\s+([^\n]+)", text)),
        "ship_date": _iso_date(_search(r"\bDELIVERY\s+(\d{1,2}/\d{1,2}/20\d{2})", text)),
        "customer_po": _search(r"CUSTOMER\s*/\s*PO#\s+([^\n]+)", text),
        "project": _search(r"Project#\s+([^\n]+)", text),
        "shipment_note": _search(r"SPECIAL REMARKS:\s*([^\n]+)", text),
        "customer": "MAXX",
    })
    result["project_no"] = _search(r"Project#\s+([A-Z0-9-]+)", text)
    line_re = re.compile(
        r"(?m)^(\d{6}(?:-\d+)?)\s+(.+?)\s+([\d,]+)\s+(Pieces?|PCS|Sets?|Set|Lot)\s+([\d,.]+)\s+([\d,.]+)\s*$",
        re.I,
    )
    for match in line_re.finditer(text):
        result["lines"].append({
            "item_code": match.group(1),
            "description": _clean(match.group(2)),
            "quantity": _number(match.group(3)),
            "unit": match.group(4),
            "unit_price": _number(match.group(5)),
            "amount": _number(match.group(6)),
        })
    known = {(line["item_code"], line["description"]) for line in result["lines"]}
    no_unit_re = re.compile(
        r"(?m)^(\d{6}(?:-\d+)?)\s+(.+?)\s+([\d,]+)\s+([\d,.]+)\s+([\d,.]+)\s*$",
        re.I,
    )
    for match in no_unit_re.finditer(text):
        key = (match.group(1), _clean(match.group(2)))
        if key in known:
            continue
        result["lines"].append({
            "item_code": match.group(1), "description": key[1],
            "quantity": _number(match.group(3)), "unit": "",
            "unit_price": _number(match.group(4)), "amount": _number(match.group(5)),
            "is_charge": match.group(1) == "510080" or "tooling" in key[1].lower(),
        })
    if not result["lines"]:
        result["warnings"].append("未从 Maxx PO 识别到产品行，请人工核对原文件。")
    if result["shipment_note"] and "ship" in result["shipment_note"].lower():
        dates = re.findall(r"\b(\d{1,2}\s+[A-Za-z]+,?\s+20\d{2})", result["shipment_note"])
        if dates:
            result["ship_date"] = _iso_date(dates[-1]) or result["ship_date"]
    return result


def _normalise_excel_header(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", _clean(value).upper())


def _right_value(values: list[Any], label_col: int) -> Any:
    for value in values[label_col + 1:]:
        if value not in (None, ""):
            return value
    return ""


def _label_value(rows: list[list[Any]], labels: tuple[str, ...]) -> Any:
    wanted = tuple(_normalise_excel_header(label) for label in labels)
    for values in rows:
        for col, value in enumerate(values):
            text = _clean(value)
            normalised = _normalise_excel_header(text)
            if not text:
                continue
            for raw_label, label in zip(labels, wanted):
                if normalised == label:
                    return _right_value(values, col)
                # SPECIAL REMARKS 等字段常把标签和值放在同一合并单元格。
                if normalised.startswith(label) and ":" in text:
                    tail = text.split(":", 1)[1].strip()
                    if tail:
                        return tail
    return ""


def _maxx_excel_lines(rows: list[list[Any]]) -> list[dict[str, Any]]:
    header_row = -1
    columns: dict[str, int] = {}
    aliases = {
        "item_code": ("ITEMNO",),
        "description": ("ITEMDESCRIPTION", "DESCRIPTION"),
        "quantity": ("QUANTITY",),
        "unit": ("UM", "UNIT"),
        "unit_price": ("USDUNITPRICE", "UNITPRICE"),
        "amount": ("USDAMOUNT", "AMOUNT"),
    }
    for row_no, values in enumerate(rows):
        current: dict[str, int] = {}
        for col, value in enumerate(values):
            token = _normalise_excel_header(value)
            for field, names in aliases.items():
                if token in names and field not in current:
                    current[field] = col
                    break
        if {"item_code", "description", "quantity"}.issubset(current):
            header_row = row_no
            columns = current
            break
    if header_row < 0:
        return []

    result: list[dict[str, Any]] = []
    for values in rows[header_row + 1:]:
        item = _clean(values[columns["item_code"]]) if columns["item_code"] < len(values) else ""
        description = _clean(values[columns["description"]]) if columns["description"] < len(values) else ""
        quantity = (
            _number(values[columns["quantity"]])
            if columns["quantity"] < len(values)
            else None
        )
        if item.upper().startswith(("TOTAL", "PAYMENT", "SPECIAL", "TERMS")):
            break
        if not item and not description and quantity is None:
            continue
        if not item or not description or quantity is None:
            continue
        unit = (
            _clean(values[columns["unit"]])
            if "unit" in columns and columns["unit"] < len(values)
            else ""
        )
        unit_price = (
            _number(values[columns["unit_price"]])
            if "unit_price" in columns and columns["unit_price"] < len(values)
            else None
        )
        amount = (
            _number(values[columns["amount"]])
            if "amount" in columns and columns["amount"] < len(values)
            else None
        )
        result.append({
            "item_code": item,
            "description": description,
            "quantity": quantity,
            "unit": unit,
            "unit_price": unit_price,
            "amount": amount,
            "is_charge": item == "510080" or "tooling" in description.lower(),
        })
    return result


def _parse_maxx_excel(rows: list[list[Any]], filename: str) -> dict[str, Any]:
    result = _base_result("maxx", filename)
    po_raw = _clean(_label_value(rows, ("P.O. NO.", "PO NO.", "PONO")))
    po_match = re.search(r"\b(PO-[A-Z0-9-]+)\b", po_raw, re.I)
    project = _clean(_label_value(rows, ("PROJECT#", "PROJECT NO.")))
    project_no = _search(r"^([A-Z0-9-]+)", project)
    delivery = _clean(_label_value(rows, ("DELIVERY MANUFACTURER CODE", "DELIVERY")))
    remarks = _clean(_label_value(rows, ("SPECIAL REMARKS",)))
    shipment_dates = re.findall(
        r"\b(?:\d{1,2}[./-]\d{1,2}[./-]20\d{2}|\d{1,2}\s+[A-Za-z]+,?\s+20\d{2})\b",
        remarks,
        re.I,
    )
    result.update({
        "po_number": (
            po_match.group(1).upper()
            if po_match
            else (po_raw.split()[0] if po_raw else "")
        ),
        "po_date": _iso_date(_label_value(rows, ("P.O. DATE", "PO DATE"))),
        "ship_date": (
            _iso_date(shipment_dates[-1])
            if shipment_dates
            else _iso_date(delivery)
        ),
        "customer_po": _clean(_label_value(rows, ("CUSTOMER/ PO#", "CUSTOMER PO#"))),
        "project": project,
        "project_no": project_no,
        "shipment_note": remarks,
        "customer": "MAXX",
        "lines": _maxx_excel_lines(rows),
    })
    if not result["lines"]:
        result["confidence"] = "low"
        result["warnings"].append(
            "已识别为 Maxx Excel，但没有读到 ITEM NO./QUANTITY 产品明细；未生成空白新单。"
        )
    return result


def _parse_shushupapa(text: str, filename: str, pages: int) -> dict[str, Any]:
    result = _base_result("shushupapa", filename, pages)
    upper = text.upper()
    compact = re.sub(r"[ \t]", "", text)
    po_number = _search(r"PO\s*Number:\s*([A-Z0-9-]+)", text)
    customer = "BJs" if "BJS CDU" in upper else (
        "Price Smart" if "PRICESMART" in upper else "SHUSHUPAPA"
    )
    result.update({
        "po_number": po_number,
        "customer_po": _search(r"PriceSmart\s*PO\s*Number:\s*([A-Z0-9-]+)", text),
        "customer": customer,
    })
    issue = _search(
        r"Issue\s*Date[^:]*:\s*([A-Za-z]+\s*\d{1,2}(?:st|nd|rd|th)?\s*,?\s*20\d{2})",
        text,
    )
    result["po_date"] = _iso_date(issue)
    delivery = _search(
        r"on\s+or\s+before\s+([A-Za-z]+\s*\d{1,2}(?:st|nd|rd|th)?\s*,?\s*20\d{2})",
        text,
    )
    result["ship_date"] = _iso_date(delivery)
    line = re.search(
        r"(?s)(\d{8})\s+(.+?)\s+([\d,]+)\s*(?:pcs)?\s+\$?([\d,.]+).*?USD\$?([\d,.]+)",
        text,
        re.I,
    )
    if line:
        item_block = text[line.start():]
        item_block = item_block.split("Total Order Value", 1)[0]
        description = _clean(line.group(2))
        if "DARTH VADER" in item_block.upper() and "DARTH VADER" not in description.upper():
            description += ", Wars, Darth Vader"
        if "LIGHT UP SABER" in item_block.upper() and "LIGHT UP SABER" not in description.upper():
            description += " Light Up Saber"
        base_line = _shushupapa_product_line(
            source_item=line.group(1),
            description=description,
            quantity=_number(line.group(3)),
            unit_price=_number(line.group(4)),
            amount=_number(line.group(5)),
            casepack=item_block,
            customer=customer,
        )
        shipments = [
            (_iso_date(match.group(1)), _number(match.group(2)))
            for match in re.finditer(
                r"(?m)^(\d{1,2}/\d{1,2}/20\d{2})\s+([\d,]+)\s*$",
                text,
            )
        ]
        if customer == "BJs" and shipments:
            result["lines"] = _split_shushupapa_shipments(
                base_line,
                po_number,
                shipments,
            )
        else:
            base_line["po_number"] = po_number
            base_line["ship_date"] = result["ship_date"]
            result["lines"].append(base_line)
    else:
        result["warnings"].append("未从 Shushupapa PO 识别到产品行，请人工核对原文件。")
    return result


def _shushupapa_header_columns(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    aliases = {
        "item_code": ("ITEM", "ITEMNO"),
        "description": ("DESCRIPTION",),
        "quantity": ("QUANTITY",),
        "unit_price": ("UNITPRICE",),
        "casepack": ("CASEPACK",),
        "amount": ("TOTALAMOUNT",),
    }
    for row_no, values in enumerate(rows):
        current: dict[str, int] = {}
        for col, value in enumerate(values):
            token = _normalise_excel_header(value)
            for field, names in aliases.items():
                if token in names and field not in current:
                    current[field] = col
        if {"item_code", "description", "quantity", "unit_price"}.issubset(current):
            return row_no, current
    return -1, {}


def _shushupapa_product_line(
    *,
    source_item: Any,
    description: Any,
    quantity: Any,
    unit_price: Any,
    amount: Any,
    casepack: Any,
    customer: str,
) -> dict[str, Any]:
    source = _clean(source_item)
    product = _clean(description)
    pack = _clean(casepack)
    upper_pack = pack.upper()
    internal_item = source
    if source == "50002008":
        internal_item = "50002008"
    elif customer == "BJs" and source in {"50002010", "50002011"}:
        # 排期首页明确：BJs 带电款使用内部货号 50002011。
        internal_item = "50002011"
    elif source in {"50002010", "50002011"}:
        internal_item = "50002010"

    factory_per_carton = None
    external_per_carton = None
    bjs_pack = re.search(
        r"(\d+)\s*PCS?\s*DARTH.*?(\d+)\s*PCS?\s*STORM",
        upper_pack,
        re.S,
    )
    price_pack = re.search(
        r"(\d+)\s*[- ]*\s*VADER.*?(\d+)\s*[- ]*\s*TROOPER",
        upper_pack,
        re.S,
    )
    pack_match = bjs_pack or price_pack
    if pack_match:
        factory_per_carton = int(pack_match.group(1))
        external_per_carton = int(pack_match.group(2))

    qty = _number(quantity)
    price = _number(unit_price)
    parsed_amount = _number(amount)
    cartons = (
        int(qty / factory_per_carton)
        if qty is not None
        and factory_per_carton
        and float(qty).is_integer()
        and int(qty) % factory_per_carton == 0
        else None
    )
    external_quantity = (
        cartons * external_per_carton
        if cartons is not None and external_per_carton is not None
        else None
    )
    total_pack = (
        (
            factory_per_carton
            if customer == "Price Smart"
            else factory_per_carton + external_per_carton
        )
        if factory_per_carton is not None and external_per_carton is not None
        else None
    )
    return {
        "item_code": internal_item,
        "po_item_code": source,
        "description": product,
        "quantity": qty,
        "external_quantity": external_quantity,
        "unit": "pcs",
        "unit_price": price,
        "amount": parsed_amount,
        "inner_pack": factory_per_carton,
        "outer_pack": total_pack,
        "cartons": cartons,
        "casepack": pack,
        "shipment_note": (
            f"与外厂混装（本厂：{factory_per_carton}PCS+外厂：{external_per_carton}PCS）"
            if factory_per_carton is not None and external_per_carton is not None
            else ""
        ),
    }


def _split_shushupapa_shipments(
    base_line: dict[str, Any],
    po_number: str,
    shipments: list[tuple[str | None, float | int | None]],
) -> list[dict[str, Any]]:
    factory_pack = _number(base_line.get("inner_pack"))
    external_pack = (
        (_number(base_line.get("outer_pack")) or 0) - (factory_pack or 0)
        if base_line.get("outer_pack") is not None
        else None
    )
    # 跟单现有排期的分单号按走货日期升序编号；同日保持 PO 原顺序。
    ordered_shipments = sorted(
        enumerate(shipments),
        key=lambda entry: (entry[1][0] or "9999-12-31", entry[0]),
    )
    result = []
    for index, (_, (ship_date, cartons)) in enumerate(ordered_shipments, 1):
        if cartons is None:
            continue
        line = dict(base_line)
        quantity = cartons * factory_pack if factory_pack is not None else None
        line.update({
            "po_number": f"{po_number}-{index}",
            "ship_date": ship_date,
            "quantity": quantity,
            "external_quantity": (
                cartons * external_pack if external_pack is not None else None
            ),
            "cartons": cartons,
            "amount": (
                round(float(quantity) * float(line["unit_price"]), 2)
                if quantity is not None and line.get("unit_price") is not None
                else None
            ),
        })
        result.append(line)
    return result


def _parse_shushupapa_excel(
    rows: list[list[Any]],
    text: str,
    filename: str,
) -> dict[str, Any]:
    result = _base_result("shushupapa", filename)
    upper = text.upper()
    customer = "BJs" if "BJS CDU" in upper else (
        "Price Smart" if "PRICESMART" in upper else "SHUSHUPAPA"
    )
    po_number = _search(r"PO\s*Number:\s*([A-Z0-9-]+)", text)
    result.update({
        "po_number": po_number,
        "customer_po": _search(
            r"PriceSmart\s*PO\s*Number:\s*([A-Z0-9-]+)",
            text,
        ),
        "po_date": _iso_date(_search(
            r"Issue\s*Date:\s*([A-Za-z]+\s*\d{1,2}(?:st|nd|rd|th)?\s*,?\s*20\d{2})",
            text,
        )),
        "ship_date": _iso_date(_search(
            r"on\s+or\s+before\s+([A-Za-z]+\s*\d{1,2}(?:st|nd|rd|th)?\s*,?\s*20\d{2})",
            text,
        )),
        "customer": customer,
    })
    header_row, columns = _shushupapa_header_columns(rows)
    if header_row < 0:
        result["confidence"] = "low"
        result["warnings"].append(
            "已识别为 Shushupapa Excel，但未找到 Item # / Quantity 产品明细表头。"
        )
        return result

    base_line = None
    for values in rows[header_row + 1:]:
        item = (
            _clean(values[columns["item_code"]])
            if columns["item_code"] < len(values)
            else ""
        )
        if not re.fullmatch(r"\d{8}", item):
            continue
        base_line = _shushupapa_product_line(
            source_item=item,
            description=(
                values[columns["description"]]
                if columns["description"] < len(values)
                else ""
            ),
            quantity=(
                values[columns["quantity"]]
                if columns["quantity"] < len(values)
                else None
            ),
            unit_price=(
                values[columns["unit_price"]]
                if columns["unit_price"] < len(values)
                else None
            ),
            amount=(
                values[columns["amount"]]
                if "amount" in columns and columns["amount"] < len(values)
                else None
            ),
            casepack=(
                values[columns["casepack"]]
                if "casepack" in columns and columns["casepack"] < len(values)
                else ""
            ),
            customer=customer,
        )
        break
    if not base_line:
        result["confidence"] = "low"
        result["warnings"].append(
            "已识别为 Shushupapa Excel，但产品明细行为空。"
        )
        return result

    shipment_header = -1
    for row_no, values in enumerate(rows):
        row_text = " ".join(_clean(value) for value in values if value not in (None, ""))
        normalized = _normalise_excel_header(row_text)
        if "CARGOREADYDATE" in normalized and "CDUUNITQUANTITY" in normalized:
            shipment_header = row_no
            break
    shipments: list[tuple[str | None, float | int | None]] = []
    if shipment_header >= 0:
        for values in rows[shipment_header + 1:]:
            nonempty = [value for value in values if value not in (None, "")]
            if len(nonempty) < 2:
                if shipments:
                    break
                continue
            ship_date = _iso_date(nonempty[0])
            cartons = _number(nonempty[1])
            if ship_date and cartons is not None:
                shipments.append((ship_date, cartons))
            elif shipments:
                break

    if customer == "BJs" and shipments:
        result["lines"] = _split_shushupapa_shipments(
            base_line,
            po_number,
            shipments,
        )
        parsed_total = sum(
            float(line.get("quantity") or 0) for line in result["lines"]
        )
        po_total = float(base_line.get("quantity") or 0)
        if abs(parsed_total - po_total) > 0.0001:
            result["confidence"] = "low"
            result["warnings"].append(
                f"BJs 分批数量合计 {parsed_total:g} 与 PO 数量 {po_total:g} 不一致，请人工核对。"
            )
    else:
        base_line["po_number"] = po_number
        base_line["ship_date"] = result["ship_date"]
        result["lines"] = [base_line]

    if customer == "BJs" and base_line.get("po_item_code") != base_line.get("item_code"):
        result["warnings"].append(
            "BJs PO 的 Item # 50002010 已按当前排期规则映射为内部货号 50002011。"
        )
    return result


_BARTER_DIGIT_TRANSLATION = str.maketrans({
    "O": "0", "Q": "0", "D": "0", "I": "1", "L": "1", "Z": "2",
    "S": "5", "$": "5", "§": "5", "G": "6", "T": "7", "B": "8",
})
_BARTER_PRODUCTS = {
    "AF QUICK DRAW": {"item_code": "88292", "product_name_zh": "绿色发声枪"},
}


def _barter_digits(value: Any) -> str:
    return re.sub(r"\D", "", _clean(value).upper().translate(_BARTER_DIGIT_TRANSLATION))


def _barter_number(value: Any) -> float | int | None:
    translated = _clean(value).upper().translate(_BARTER_DIGIT_TRANSLATION)
    return _number(translated)


def _barter_contract_range(filename: str, page_count: int) -> tuple[int, int] | None:
    match = re.search(r"TNT\s*(\d{5})\s*[-_]\s*(\d{5})", Path(filename).stem, re.I)
    if not match:
        return None
    start, end = int(match.group(1)), int(match.group(2))
    return (start, end) if end >= start and end - start + 1 == page_count else None


def _barter_customer_po(text: str) -> str:
    matches: list[str] = []
    for line in text.splitlines():
        upper = line.upper()
        if "CUSTOMER" not in upper:
            continue
        tail = upper.split("USA", 1)[-1]
        matches.extend(re.findall(r"\d{10}", _barter_digits(tail)))
    return next(
        (value for value in matches if value.startswith("0990")),
        next((value for value in matches if value.startswith("09")), matches[0] if matches else ""),
    )


def _barter_date_code(text: str, item_code: str) -> str:
    compact = re.sub(r"\s+", "", text.upper()).translate(_BARTER_DIGIT_TRANSLATION)
    match = re.search(rf"WM{re.escape(item_code)}RR[-~]?([0-9]{{4}})", compact)
    if not match:
        match = re.search(r"DATEC0DE.*?RR[-~]?([0-9]{4})", compact)
    return f"WM{item_code}RR-{match.group(1)}" if match else ""


def _parse_barter_page(
    text: str,
    *,
    page_number: int,
    expected_contract: str | None,
    ocr_used: bool,
) -> tuple[dict[str, Any], list[str]]:
    warnings: list[str] = []
    contract_match = re.search(r"CONTRACT\s*[^\n]*?TNT\s*/?\s*F?\s*([0-9OQDISZTBG]{5})", text, re.I)
    parsed_contract = (
        f"TNT/F{_barter_digits(contract_match.group(1))}"
        if contract_match and len(_barter_digits(contract_match.group(1))) == 5
        else ""
    )
    contract = expected_contract or parsed_contract
    if expected_contract and parsed_contract and expected_contract != parsed_contract:
        warnings.append(f"第 {page_number} 页合同号 OCR 为 {parsed_contract}，已按文件名页序校正为 {expected_contract}")

    delivery = _search(
        r"DELIVERY\s*[:<#]?\s*([A-Z]{3,9}[\s.\-]*\d{1,2}\s*,?\s*20\d{2})",
        text,
    )
    ship_date = _iso_date(delivery)
    customer_po = _barter_customer_po(text)

    product_line = next((
        _clean(line) for line in text.splitlines()
        if re.search(r"\bP[CO]S\b", line, re.I) and re.search(r"/\s*P[CO]S", line, re.I)
    ), "")
    product_match = re.search(
        r"^\s*([A-Z0-9]{4,})\s+(.+?)\s+([0-9OQDISZTBG§$,]+)\s+P[CO]S\s+"
        r"([0-9OQDISZTBG§$,.]+)\s*/\s*P[CO]S[,]?\s+([0-9OQDISZTBG§$,. ]+)",
        product_line,
        re.I,
    )
    parsed_item = _barter_digits(product_match.group(1)) if product_match else ""
    description = _clean(product_match.group(2)).upper() if product_match else ""
    known_product = next((
        (name, values) for name, values in _BARTER_PRODUCTS.items()
        if name in description or name in text.upper()
    ), None)
    if known_product:
        description = known_product[0]
        item_code = known_product[1]["item_code"]
        product_name_zh = known_product[1]["product_name_zh"]
    else:
        item_code = parsed_item
        product_name_zh = description

    parsed_quantity = _barter_number(product_match.group(3)) if product_match else None
    unit_price = _barter_number(product_match.group(4)) if product_match else None
    parsed_amount = _barter_number(product_match.group(5)) if product_match else None
    carton_match = re.search(r"\(\s*([0-9OQDISZTBG§$, ]+)\s+[CG]TN\s*\)", text, re.I)
    cartons = _barter_number(carton_match.group(1)) if carton_match else None
    pack_match = re.search(r"([0-9OQDISZTBG§$]+)\s+P[CO]\s+PER\s+MASTER", text, re.I)
    outer_pack = _barter_number(pack_match.group(1)) if pack_match else None
    quantity = parsed_quantity
    if cartons and outer_pack:
        calculated_quantity = int(float(cartons) * float(outer_pack))
        if parsed_quantity is None or abs(float(parsed_quantity) - calculated_quantity) > 0.001:
            warnings.append(f"第 {page_number} 页数量已按箱数×装箱数校正为 {calculated_quantity}")
        quantity = calculated_quantity
    if quantity and parsed_amount:
        implied_price = round(float(parsed_amount) / float(quantity), 4)
        if (
            0 < implied_price < 1000
            and (
                unit_price is None
                or abs(float(unit_price) * float(quantity) - float(parsed_amount)) > 0.05
            )
        ):
            warnings.append(f"第 {page_number} 页 HKD 单价已按行金额÷数量校正为 {implied_price:g}")
            unit_price = implied_price
    if unit_price is not None and not 0 < float(unit_price) < 1000:
        warnings.append(f"第 {page_number} 页 HKD 单价 OCR 异常，已留待历史排期或人工校正")
        unit_price = None
    amount = parsed_amount
    if quantity is not None and unit_price is not None:
        calculated_amount = round(float(quantity) * float(unit_price), 2)
        if parsed_amount is None or abs(float(parsed_amount) - calculated_amount) > 0.05:
            warnings.append(f"第 {page_number} 页金额已按数量×HKD单价校正为 {calculated_amount:.2f}")
        amount = calculated_amount

    line = {
        "po_number": contract,
        "customer_po": customer_po,
        "ship_date": ship_date,
        "item_code": item_code,
        "product_name_zh": product_name_zh,
        "description": description,
        "quantity": quantity,
        "unit": "pcs",
        "unit_price": unit_price,
        "amount": amount,
        "outer_pack": outer_pack,
        "cartons": cartons,
        "date_code": _barter_date_code(text, item_code) if item_code else "",
        "color_box": "英文盒（2026版本）" if "2026" in text and "GREEN BOX" in text.upper() else "",
        "customer_label": "RFID" if "RFID" in text.upper() else "",
        "label": "外箱利宝" if "CARTON LABEL" in text.upper() else "",
        "battery": "AA（2pcs）" if re.search(r"2\s+P[CO]S\s+.*AA", text, re.I) else "",
        "country": "WM USA",
        "source_page": page_number,
        "_ocr_used": ocr_used,
        "_ocr_reconciled": bool(warnings),
    }
    return line, warnings


def _parse_barter_pages(
    page_texts: list[str], filename: str, *, ocr_used: bool = False,
) -> dict[str, Any]:
    result = _base_result("barter", filename, len(page_texts))
    result["customer"] = "WM USA"
    contract_range = _barter_contract_range(filename, len(page_texts))
    page_warnings: list[str] = []
    for page_number, text in enumerate(page_texts, start=1):
        expected = f"TNT/F{contract_range[0] + page_number - 1}" if contract_range else None
        line, warnings = _parse_barter_page(
            text,
            page_number=page_number,
            expected_contract=expected,
            ocr_used=ocr_used,
        )
        page_warnings.extend(warnings)
        if line.get("po_number") or any(line.get(key) for key in ("item_code", "quantity", "customer_po")):
            result["lines"].append(line)
    if result["lines"]:
        first = result["lines"][0]
        result.update({
            "po_number": first.get("po_number") or "",
            "customer_po": first.get("customer_po") or "",
            "ship_date": first.get("ship_date"),
        })
    if ocr_used:
        result["confidence"] = "medium"
        result["warnings"].append(
            f"该 Barter 合同为扫描件，已逐页 OCR 提取 {len(page_texts)} 页；导出前请复核合同号、客户 PO、数量、日期码及 CRD。"
        )
    if contract_range:
        result["warnings"].append(
            f"文件名合同范围 TNT/F{contract_range[0]}–TNT/F{contract_range[1]} 与 {len(page_texts)} 页逐页对应。"
        )
    if page_warnings:
        result["warnings"].append("；".join(page_warnings[:8]) + ("；其余校正详见逐行预览" if len(page_warnings) > 8 else ""))
    if not result["lines"]:
        result["confidence"] = "low"
        result["warnings"].append("未识别到 Barter 产品行，请核对合同版式。")
    return result


def _parse_barter(text: str, filename: str, pages: int) -> dict[str, Any]:
    """Backward-compatible entry point for native-text Barter documents."""
    return _parse_barter_pages(text.split("\f") if "\f" in text else [text], filename)


def _parse_excel(source: str | Path | bytes | BinaryIO, filename: str) -> dict[str, Any]:
    payload = io.BytesIO(source) if isinstance(source, bytes) else source
    wb = load_workbook(payload, data_only=True, read_only=True)
    try:
        text_parts: list[str] = []
        rows: list[list[Any]] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                values = list(row)
                rows.append(values)
                text_parts.extend(_clean(v) for v in values if v not in (None, ""))
        text = "\n".join(text_parts)
        client = detect_client(text, filename)
        if client == "maxx":
            return _parse_maxx_excel(rows, filename)
        if client == "shushupapa":
            return _parse_shushupapa_excel(rows, text, filename)
        result = _base_result(client, filename)
        result["warnings"].append("Excel PO 已读取；字段版式不固定，导入前请在待确认表核对。")
        for values in rows:
            joined = " | ".join(_clean(v) for v in values if v not in (None, ""))
            if not result["po_number"]:
                result["po_number"] = _search(r"(?:P\.?O\.?\s*(?:NO\.?)?|PURCHASE ORDER NO\.?)\s*[:#]?\s*([A-Z0-9-]+)", joined)
        return result
    finally:
        wb.close()


def _check_line_math(result: dict[str, Any]) -> None:
    """自洽校验：逐行 数量×单价 = 行金额（所有客户统一走一遍）。"""
    for line in result.get("lines") or []:
        qty = _number(line.get("quantity")) or 0
        price = _number(line.get("unit_price")) or 0
        amount = _number(line.get("amount")) or 0
        if qty > 0 and price > 0 and amount > 0:
            expect = round(qty * price, 2)
            if abs(expect - amount) > 0.05:
                result.setdefault("warnings", []).append(
                    f"{line.get('item_code') or line.get('description', '')[:20]}: "
                    f"数量×单价({qty}×{price}={expect})与金额({amount})不符，请人工核对（注意是否有折扣）"
                )


def parse_po(source: str | Path | bytes | BinaryIO, filename: str = "") -> dict[str, Any]:
    name = filename or (Path(source).name if isinstance(source, (str, Path)) else "upload")
    if name.lower().endswith((".xlsx", ".xlsm")):
        result = _parse_excel(source, name)
        _check_line_math(result)
        return result
    page_texts, ocr_used = read_pdf_pages(source)
    text = "\n".join(page_texts)
    pages = len(page_texts)
    client = detect_client(text, name)
    if client == "maxx":
        result = _parse_maxx(text, name, pages)
    elif client == "shushupapa":
        result = _parse_shushupapa(text, name, pages)
    elif client == "barter":
        result = _parse_barter_pages(page_texts, name, ocr_used=ocr_used)
    else:
        result = _base_result("", name, pages)
        result["confidence"] = "low"
        result["warnings"].append("无法自动识别客户，请在页面选择客户后再处理。")
    if not result.get("po_number"):
        result["warnings"].append("未识别到订单号。")
    _check_line_math(result)
    return result
