# -*- coding: utf-8 -*-
"""Generate customer-order workbooks from the corresponding schedule template."""
from __future__ import annotations

import re
from copy import copy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO, Callable, Collection, Iterable, Mapping, Sequence

import openpyxl
import xlrd
from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange

from app.services.legacy_excel_bridge import (
    convert_legacy_xls_to_xlsx,
    is_legacy_xls_workbook,
)


def _normalise(value: Any) -> str:
    if value in (None, ""):
        return ""
    text = str(value).strip().lower()
    return re.sub(r"[\s\u3000_／/()（）【】\[\]：:·.,，。-]+", "", text)


def _copy_style(source, target) -> None:
    if not source.has_style:
        return
    target._style = copy(source._style)


_SHEET_REFERENCE = re.compile(
    r"(?P<sheet>'(?:[^']|'')+'|[A-Za-z0-9_\u4e00-\u9fff]+)!"
    r"(?P<cell>\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)"
)


def _safe_header_value(value: Any, target_sheet: str) -> Any:
    """Remove references to sheets that are intentionally absent from new-order files."""
    if not isinstance(value, str) or not value.startswith("="):
        return value

    target_key = target_sheet.replace("'", "").strip().casefold()

    def replace(match: re.Match[str]) -> str:
        sheet = match.group("sheet").strip("'").replace("''", "'").strip().casefold()
        return match.group(0) if sheet == target_key else "0"

    return _SHEET_REFERENCE.sub(replace, value)


def _xls_colour(book, colour_index: int | None, default: str = "000000") -> str:
    if colour_index is None:
        return default
    rgb = getattr(book, "colour_map", {}).get(colour_index)
    return default if not rgb else "".join(f"{part:02X}" for part in rgb)


def _load_xls(
    path_or_bytes: str | Path | bytes,
    sheet_names: Sequence[str] = (),
) -> Workbook:
    kwargs = {"formatting_info": True}
    if isinstance(path_or_bytes, bytes):
        book = xlrd.open_workbook(file_contents=path_or_bytes, **kwargs)
    else:
        book = xlrd.open_workbook(str(path_or_bytes), **kwargs)
    result = Workbook()
    result.remove(result.active)
    selected = [sheet for sheet in book.sheets() if not sheet_names or sheet.name in sheet_names]
    if not selected:
        selected = book.sheets()
    for source in selected:
        ws = result.create_sheet((source.name or "Sheet")[:31])
        for row_index in range(source.nrows):
            row_info = source.rowinfo_map.get(row_index)
            if row_info and row_info.height:
                ws.row_dimensions[row_index + 1].height = row_info.height / 20
            for col_index in range(source.ncols):
                old = source.cell(row_index, col_index)
                value = old.value
                if old.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                    except (TypeError, ValueError):
                        pass
                cell = ws.cell(row_index + 1, col_index + 1, value)
                try:
                    xf = book.xf_list[old.xf_index]
                    font = book.font_list[xf.font_index]
                    cell.font = Font(
                        name=font.name or "Arial",
                        size=(font.height or 200) / 20,
                        bold=bool(font.bold),
                        italic=bool(font.italic),
                        underline="single" if font.underline_type else None,
                        color=_xls_colour(book, font.colour_index),
                    )
                    background = book.colour_map.get(xf.background.pattern_colour_index)
                    if background:
                        cell.fill = PatternFill(
                            "solid",
                            fgColor="".join(f"{part:02X}" for part in background),
                        )
                    border = xf.border
                    side = lambda style, colour: Side(
                        style="thin" if style else None,
                        color=_xls_colour(book, colour),
                    )
                    cell.border = Border(
                        left=side(border.left_line_style, border.left_colour_index),
                        right=side(border.right_line_style, border.right_colour_index),
                        top=side(border.top_line_style, border.top_colour_index),
                        bottom=side(border.bottom_line_style, border.bottom_colour_index),
                    )
                    alignment = xf.alignment
                    horizontal = {
                        1: "left", 2: "center", 3: "right", 4: "fill",
                        5: "justify", 6: "centerContinuous", 7: "distributed",
                    }.get(alignment.hor_align)
                    vertical = {0: "top", 1: "center", 2: "bottom", 3: "justify", 4: "distributed"}.get(
                        alignment.vert_align
                    )
                    cell.alignment = Alignment(
                        horizontal=horizontal,
                        vertical=vertical,
                        wrap_text=bool(alignment.text_wrapped),
                    )
                    cell.number_format = book.format_map[xf.format_key].format_str or "General"
                    cell.protection = Protection(
                        locked=bool(xf.protection.cell_locked),
                        hidden=bool(xf.protection.formula_hidden),
                    )
                except (AttributeError, IndexError, KeyError, TypeError):
                    pass
        for col_index, info in source.colinfo_map.items():
            if info.width:
                ws.column_dimensions[get_column_letter(col_index + 1)].width = max(1, info.width / 256)
        for row_low, row_high, col_low, col_high in source.merged_cells:
            ws.merge_cells(
                start_row=row_low + 1,
                end_row=row_high,
                start_column=col_low + 1,
                end_column=col_high,
            )
    return result


def load_workbook_compatible(
    source: str | Path | bytes | BinaryIO,
    *,
    filename: str = "",
    data_only: bool = False,
    sheet_names: Sequence[str] = (),
) -> Workbook:
    """Load xlsx/xlsm/xls from a path, bytes, or file-like object."""
    suffix = Path(filename or str(source) if not isinstance(source, bytes) else filename).suffix.lower()
    if suffix == ".xls":
        if hasattr(source, "read"):
            source = source.read()
        return _load_xls(source, sheet_names)
    workbook_options = {
        "data_only": data_only,
        "keep_links": True,
        "keep_vba": suffix in {".xlsm", ".xltm"},
    }
    if hasattr(source, "read"):
        payload = source.read()
        return openpyxl.load_workbook(BytesIO(payload), **workbook_options)
    if isinstance(source, bytes):
        return openpyxl.load_workbook(BytesIO(source), **workbook_options)
    return openpyxl.load_workbook(source, **workbook_options)


def load_complete_workbook_compatible(
    source: str | Path | bytes | BinaryIO,
    *,
    filename: str = "",
) -> Workbook:
    """Load a workbook for a full-copy export.

    Legacy ``.xls`` files must be converted by Excel/LibreOffice before editing;
    the lightweight xlrd reconstruction used by standalone exports cannot retain
    drawings, formulas, print settings, or workbook-level metadata.
    """
    suffix = Path(filename).suffix.lower()
    if suffix != ".xls":
        return load_workbook_compatible(source, filename=filename)
    if hasattr(source, "read"):
        payload = source.read()
    elif isinstance(source, bytes):
        payload = source
    else:
        payload = Path(source).read_bytes()
    if not is_legacy_xls_workbook(payload):
        raise ValueError("上传文件扩展名为 .xls，但内容不是有效的旧版 Excel 工作簿")
    converted = convert_legacy_xls_to_xlsx(payload)
    return openpyxl.load_workbook(BytesIO(converted), data_only=False, keep_links=True)


def _alias_lookup(field_aliases: Mapping[str, Sequence[str]]) -> dict[str, str]:
    lookup: dict[str, str] = {}
    for field, aliases in field_aliases.items():
        for alias in (field, *aliases):
            normalised = _normalise(alias)
            if normalised:
                lookup[normalised] = field
    return lookup


def detect_header(
    ws,
    field_aliases: Mapping[str, Sequence[str]],
    *,
    scan_rows: int = 20,
) -> tuple[int, dict[int, str]]:
    """Return the highest-scoring header row and its column-to-field map."""
    aliases = _alias_lookup(field_aliases)
    best_row = 1
    best_map: dict[int, str] = {}
    for row in range(1, min(ws.max_row or 1, scan_rows) + 1):
        current: dict[int, str] = {}
        for col in range(1, (ws.max_column or 1) + 1):
            value = _normalise(ws.cell(row, col).value)
            if not value:
                continue
            field = aliases.get(value)
            if not field:
                field = next(
                    (
                        candidate
                        for alias, candidate in aliases.items()
                        if len(alias) >= 2 and (alias in value or value in alias)
                    ),
                    "",
                )
            if field and field not in current.values():
                current[col] = field
        if len(current) > len(best_map):
            best_row, best_map = row, current
    return best_row, best_map


def select_sheet(
    workbook: Workbook,
    field_aliases: Mapping[str, Sequence[str]],
    sheet_names: Sequence[str] = (),
) -> tuple[Any, int, dict[int, str]]:
    for name in sheet_names:
        if name in workbook.sheetnames:
            ws = workbook[name]
            header_row, column_map = detect_header(ws, field_aliases)
            return ws, header_row, column_map
    best = None
    for ws in workbook.worksheets:
        header_row, column_map = detect_header(ws, field_aliases)
        candidate = (len(column_map), ws, header_row, column_map)
        if best is None or candidate[0] > best[0]:
            best = candidate
    if best is None:
        raise ValueError("模板中没有可用工作表")
    return best[1], best[2], best[3]


def _copy_headers(source, target, header_rows: int, max_col: int) -> None:
    for row in range(1, header_rows + 1):
        if source.row_dimensions[row].height:
            target.row_dimensions[row].height = source.row_dimensions[row].height
        for col in range(1, max_col + 1):
            old = source.cell(row, col)
            new = target.cell(
                row,
                col,
                _safe_header_value(old.value, target.title),
            )
            _copy_style(old, new)
    for col in range(1, max_col + 1):
        letter = get_column_letter(col)
        dimension = source.column_dimensions[letter]
        target.column_dimensions[letter].width = dimension.width
        target.column_dimensions[letter].hidden = dimension.hidden
    for merged in source.merged_cells.ranges:
        if merged.max_row <= header_rows and merged.max_col <= max_col:
            target.merge_cells(str(merged))
    target.freeze_panes = source.freeze_panes or f"A{header_rows + 1}"
    target.sheet_view.showGridLines = source.sheet_view.showGridLines


def _style_source_row(ws, header_rows: int, max_col: int, requested: int | None) -> int:
    if requested and requested <= ws.max_row:
        return requested
    candidates = range(header_rows + 1, min(ws.max_row or header_rows + 1, header_rows + 30) + 1)
    return max(
        candidates,
        key=lambda row: sum(1 for col in range(1, max_col + 1) if ws.cell(row, col).has_style),
        default=header_rows,
    )


def _safe_value(value: Any) -> Any:
    if isinstance(value, date) and not isinstance(value, datetime):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, (dict, list, tuple, set)):
        return " / ".join(str(item) for item in value)
    return value


_DATE_FIELD_NAMES = {
    "order_date", "po_date", "change_date", "inspection_date", "customer_inspection_date",
    "factory_inspection", "customer_inspection", "complete_date", "finish_date", "ship_date",
    "po_ship_date", "requested_ship_date", "fcd", "factory_commit_date", "line_start_date",
    "booking_date", "warehouse_date", "material_ready_date", "plastic_ready_date",
    "carton_ready_date", "injection_ready_date", "pre_chase_material_date", "revision_date",
}


def _safe_record_value(field: str, value: Any) -> Any:
    if field == "date_code" or field not in _DATE_FIELD_NAMES:
        return _safe_value(value)
    if isinstance(value, (date, datetime)) or value in (None, ""):
        return _safe_value(value)
    text = str(value).strip()
    for format_string in (
        "%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%d-%b-%Y", "%d-%B-%Y",
        "%m/%d/%Y", "%d/%m/%Y",
    ):
        try:
            return datetime.strptime(text, format_string)
        except ValueError:
            continue
    return _safe_value(value)


ITEM_FIELD_CANDIDATES = ("item_no", "item", "sku", "product_no", "material_no")
PRODUCT_FIELD_CANDIDATES = (
    "product_name",
    "product",
    "description",
    "english_name",
)


def _item_key(value: Any) -> str:
    if value in (None, ""):
        return ""
    text = str(value).strip().upper()
    if text.endswith(".0"):
        text = text[:-2]
    return re.sub(r"[^A-Z0-9]", "", text)


def _display_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def inherit_product_names(
    source,
    header_row: int,
    column_map: Mapping[int, str],
    rows: list[dict[str, Any]],
) -> dict[str, Any]:
    """Use the schedule's unique item-to-name mapping as the product master.

    A source item may occur on many historical rows.  The name is inherited only
    when those rows resolve to one non-empty normalized name.  Conflicting source
    names are reported and never guessed.  A PO-provided English description is
    preserved; schedule English names only fill blanks.
    """
    field_columns = {field: col for col, field in column_map.items()}
    item_field = next((field for field in ITEM_FIELD_CANDIDATES if field in field_columns), "")
    product_fields = [
        field for field in PRODUCT_FIELD_CANDIDATES if field in field_columns
    ]
    if not item_field or not product_fields:
        return {
            "item_field": item_field,
            "product_fields": product_fields,
            "mapped_items": 0,
            "applied": 0,
            "conflicts": [],
        }

    source_values: dict[str, dict[str, dict[str, str]]] = {}
    item_col = field_columns[item_field]
    for row_no in range(header_row + 1, (source.max_row or header_row) + 1):
        key = _item_key(source.cell(row_no, item_col).value)
        if not key:
            continue
        per_item = source_values.setdefault(key, {})
        for field in product_fields:
            value = _display_text(source.cell(row_no, field_columns[field]).value)
            if not value or value.startswith("="):
                continue
            normalized = _normalise(value)
            if normalized:
                per_item.setdefault(field, {})[normalized] = value

    unique_map: dict[str, dict[str, str]] = {}
    conflicts: list[dict[str, Any]] = []
    for item, fields in source_values.items():
        for field, values in fields.items():
            if len(values) == 1:
                unique_map.setdefault(item, {})[field] = next(iter(values.values()))
            elif len(values) > 1:
                conflicts.append({
                    "item": item,
                    "field": field,
                    "values": list(values.values())[:8],
                })

    applied = 0
    inherited_rows: list[dict[str, str]] = []
    for record in rows:
        key = _item_key(record.get(item_field))
        inherited = unique_map.get(key, {})
        if not inherited:
            continue
        for field, value in inherited.items():
            old = _display_text(record.get(field))
            if _normalise(old) == _normalise(value):
                continue
            # 英文品名在有 PO DESCRIPTION 时以本次 PO 为准；排期只用于补空。
            # 中文“产品名称”仍按用户要求由货号对应的排期产品主数据覆盖。
            if field == "english_name" and old:
                continue
            record[field] = value
            if field in {"product_name", "product"} and "product_name_source" in record:
                record["product_name_source"] = "最新排期按货号唯一继承"
            applied += 1
            inherited_rows.append({
                "item": key,
                "field": field,
                "po_value": old,
                "schedule_value": value,
            })

    return {
        "item_field": item_field,
        "product_fields": product_fields,
        "mapped_items": len(unique_map),
        "applied": applied,
        "inherited_rows": inherited_rows,
        "conflicts": conflicts,
        "rule": "same_schedule_same_item_unique_name",
    }


_TOTAL_MARKERS = {_normalise(value) for value in ("合计", "總計", "总计", "小计", "小計")}
_MAX_EXCEL_ROW = 1_048_576
_FORMULA_RANGE = re.compile(
    r"^(?:(?P<sheet>'(?:[^']|'')+'|[^!]+)!)?"
    r"(?P<start>\$?[A-Z]{1,3}\$?\d+)"
    r"(?::(?P<end>\$?[A-Z]{1,3}\$?\d+))?$"
)
_FORMULA_COORD = re.compile(r"^(?P<column>\$?[A-Z]{1,3})(?P<row_abs>\$?)(?P<row>\d+)$")


_DETAIL_FIELDS = {
    "po_no", "po_number", "customer_po", "contract", "contract_no", "huaxing_po",
    "production_no", "customer_release_no", "item", "item_no", "item_full", "sku",
    "product_no", "material_no", "quantity", "qty",
}


def _row_has_total_marker(ws, row_no: int, max_col: int) -> bool:
    return any(
        isinstance(ws.cell(row_no, col_no).value, str)
        and not ws.cell(row_no, col_no).value.startswith("=")
        and _normalise(ws.cell(row_no, col_no).value) in _TOTAL_MARKERS
        for col_no in range(1, max_col + 1)
    )


def _nearest_total_row(ws, header_row: int, before_row: int, max_col: int) -> int | None:
    return next(
        (
            row_no
            for row_no in range(min(before_row, ws.max_row or before_row), header_row, -1)
            if _row_has_total_marker(ws, row_no, max_col)
        ),
        None,
    )


def _find_append_row(
    ws,
    header_row: int,
    max_col: int,
    *,
    detail_columns: Sequence[int] = (),
) -> int:
    """Find the detail boundary before totals, subtotals, or summary sections."""
    if detail_columns:
        merged_rows = {
            row_no
            for merged in ws.merged_cells.ranges
            if merged.max_col > merged.min_col
            for row_no in range(merged.min_row, merged.max_row + 1)
        }
        detail_rows = []
        for row_no in range(header_row + 1, (ws.max_row or header_row) + 1):
            if row_no in merged_rows:
                continue
            values = [ws.cell(row_no, col_no).value for col_no in detail_columns]
            if any(
                isinstance(value, str) and not value.startswith("=")
                and _normalise(value) in _TOTAL_MARKERS
                for value in values
            ):
                continue
            if any(
                value not in (None, "")
                and not (
                    isinstance(value, str)
                    and (value.startswith("=") or not value.strip())
                )
                for value in values
            ):
                detail_rows.append(row_no)
        if detail_rows:
            return detail_rows[-1] + 1

    # Header-only templates and unusual layouts fall back to the first explicit
    # total marker, which is the established boundary in the 360/Yinhui sheets.
    for row_no in range(header_row + 1, (ws.max_row or header_row) + 1):
        for col_no in range(1, max_col + 1):
            value = ws.cell(row_no, col_no).value
            if not isinstance(value, str) or value.startswith("="):
                continue
            if _normalise(value) in _TOTAL_MARKERS:
                return row_no
    populated = [
        row_no
        for row_no in range(header_row + 1, (ws.max_row or header_row) + 1)
        if any(ws.cell(row_no, col_no).value not in (None, "") for col_no in range(1, max_col + 1))
    ]
    return (populated[-1] + 1) if populated else header_row + 1


def _first_total_section_rows(
    ws,
    header_row: int,
    max_col: int,
    *,
    detail_columns: Sequence[int] = (),
) -> tuple[int, int, int]:
    """Return append row, first-total row, and reusable blank-row count.

    Some customer schedules contain several independent order sections on the
    same worksheet.  New orders belong only to the first section: reuse the
    reserved blank rows after its last real order and insert only the shortfall
    immediately before that section's first total row.
    """
    total_row = next(
        (
            row_no
            for row_no in range(header_row + 1, (ws.max_row or header_row) + 1)
            if any(
                isinstance(ws.cell(row_no, col_no).value, str)
                and not ws.cell(row_no, col_no).value.startswith("=")
                and _normalise(ws.cell(row_no, col_no).value) in _TOTAL_MARKERS
                for col_no in range(1, max_col + 1)
            )
        ),
        None,
    )
    if total_row is None:
        raise ValueError(f"{ws.title} 未找到首段合计行，无法安全确定新单写入区域")

    inspected_columns = tuple(detail_columns) or tuple(range(1, max_col + 1))
    last_detail_row = next(
        (
            row_no
            for row_no in range(total_row - 1, header_row, -1)
            if any(
                ws.cell(row_no, col_no).value not in (None, "")
                and not (
                    isinstance(ws.cell(row_no, col_no).value, str)
                    and ws.cell(row_no, col_no).value.startswith("=")
                )
                for col_no in inspected_columns
            )
        ),
        header_row,
    )
    append_row = last_detail_row + 1
    reusable_rows = 0
    for row_no in range(append_row, total_row):
        if any(ws.cell(row_no, col_no).value not in (None, "") for col_no in range(1, max_col + 1)):
            break
        reusable_rows += 1
    return append_row, total_row, reusable_rows


def _representative_data_row(
    ws,
    header_row: int,
    append_row: int,
    max_col: int,
    column_map: Mapping[int, str],
) -> int:
    """Use the nearest preceding detail row, skipping totals and group headings."""
    return _nearest_detail_row(
        ws,
        header_row,
        append_row,
        max_col,
        column_map=column_map,
    )


def _nearest_detail_row(
    ws,
    header_row: int,
    append_row: int,
    max_col: int,
    *,
    column_map: Mapping[int, str] | None = None,
) -> int:
    """Find the closest real order-detail row above an insertion boundary."""
    merged_rows = {
        row_no
        for merged in ws.merged_cells.ranges
        if merged.max_col > merged.min_col
        for row_no in range(merged.min_row, merged.max_row + 1)
    }
    mapped_columns = tuple((column_map or {}).keys())
    for row_no in range(append_row - 1, header_row, -1):
        values = [ws.cell(row_no, col_no).value for col_no in range(1, max_col + 1)]
        text_values = [value for value in values if isinstance(value, str) and not value.startswith("=")]
        if any(_normalise(value) in _TOTAL_MARKERS for value in text_values):
            continue
        if row_no in merged_rows:
            continue
        if mapped_columns:
            if any(ws.cell(row_no, col_no).value not in (None, "") for col_no in mapped_columns):
                return row_no
            continue
        if any(value not in (None, "") and not (isinstance(value, str) and value.startswith("=")) for value in values):
            return row_no
    return max(header_row, append_row - 1)


def _nearest_group_title_row(
    ws,
    header_row: int,
    detail_style_row: int,
    max_col: int,
) -> tuple[int, list[CellRange]]:
    """Find the nearest horizontally merged item/title row above a detail row."""
    for row_no in range(detail_style_row - 1, header_row, -1):
        row_merges = [
            CellRange(str(merged))
            for merged in ws.merged_cells.ranges
            if merged.min_row == row_no
            and merged.max_row == row_no
            and merged.max_col > merged.min_col
            and merged.max_col <= max_col
        ]
        if not row_merges:
            continue
        values = [ws.cell(row_no, col_no).value for col_no in range(1, max_col + 1)]
        text_values = [
            value
            for value in values
            if isinstance(value, str) and not value.startswith("=")
        ]
        if any(_normalise(value) in _TOTAL_MARKERS for value in text_values):
            continue
        if any(value not in (None, "") for value in values):
            return row_no, row_merges
    raise ValueError("当前排期没有可复用的货号/产品标题行格式")


def _shift_cell_range(
    value: CellRange,
    insert_row: int,
    amount: int,
    *,
    expand_adjacent: bool = False,
) -> CellRange:
    shifted = CellRange(str(value))
    if shifted.min_row >= insert_row:
        shifted.min_row = min(shifted.min_row + amount, _MAX_EXCEL_ROW)
        shifted.max_row = min(shifted.max_row + amount, _MAX_EXCEL_ROW)
    elif shifted.max_row >= insert_row or (expand_adjacent and shifted.max_row == insert_row - 1):
        shifted.max_row = min(shifted.max_row + amount, _MAX_EXCEL_ROW)
    return shifted


def _shift_multi_range(value: MultiCellRange, insert_row: int, amount: int) -> MultiCellRange:
    shifted = MultiCellRange()
    for cell_range in value.ranges:
        shifted.add(str(_shift_cell_range(cell_range, insert_row, amount)))
    return shifted


def _formula_sheet_name(value: str | None) -> str:
    if not value:
        return ""
    unquoted = value[1:-1].replace("''", "'") if value.startswith("'") and value.endswith("'") else value
    return unquoted.strip()


def _formula_coordinate(value: str) -> tuple[str, bool, int] | None:
    match = _FORMULA_COORD.fullmatch(value)
    if not match:
        return None
    return match.group("column"), bool(match.group("row_abs")), int(match.group("row"))


def _formula_coordinate_with_row(value: str, row_no: int) -> str:
    parsed = _formula_coordinate(value)
    if not parsed:
        return value
    column, absolute, _ = parsed
    return f"{column}{'$' if absolute else ''}{row_no}"


def _rewrite_formula_for_insert(
    formula: str,
    *,
    formula_sheet: str,
    target_sheet: str,
    formula_row: int,
    insert_row: int,
    amount: int,
) -> str:
    """Apply Excel-like row insertion updates to references targeting the edited sheet."""
    from openpyxl.formula.tokenizer import Tokenizer

    tokenizer = Tokenizer(formula)
    for token in tokenizer.items:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        match = _FORMULA_RANGE.fullmatch(token.value)
        if not match:
            continue
        referenced_sheet = _formula_sheet_name(match.group("sheet"))
        if referenced_sheet:
            if referenced_sheet.casefold() != target_sheet.casefold():
                continue
        elif formula_sheet.casefold() != target_sheet.casefold():
            continue
        start_value = match.group("start")
        end_value = match.group("end")
        start = _formula_coordinate(start_value)
        end = _formula_coordinate(end_value) if end_value else None
        if not start:
            continue
        start_row = start[2]
        if end:
            end_row = end[2]
            if start_row >= insert_row:
                start_row += amount
                end_row += amount
            elif end_row >= insert_row:
                end_row += amount
            elif formula_sheet.casefold() == target_sheet.casefold() and formula_row >= insert_row and end_row == insert_row - 1:
                # A total formula directly below the detail range expands with the inserted rows.
                end_row += amount
            start_value = _formula_coordinate_with_row(start_value, start_row)
            end_value = _formula_coordinate_with_row(end_value, end_row)
        elif start_row >= insert_row:
            start_value = _formula_coordinate_with_row(start_value, start_row + amount)
        prefix = f"{match.group('sheet')}!" if match.group("sheet") else ""
        token.value = prefix + start_value + (f":{end_value}" if end_value else "")
    return tokenizer.render()


def _extend_subtotal_to_previous_row(formula: str, target_row: int) -> str:
    """Make a moved SUBTOTAL row include every newly inserted detail row above it."""
    from openpyxl.formula.tokenizer import Tokenizer

    if "SUBTOTAL" not in formula.upper():
        return formula
    tokenizer = Tokenizer(formula)
    for token in tokenizer.items:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        match = _FORMULA_RANGE.fullmatch(token.value)
        if not match or match.group("sheet") or not match.group("end"):
            continue
        start = _formula_coordinate(match.group("start"))
        end = _formula_coordinate(match.group("end"))
        if not start or not end or start[2] >= target_row or end[2] >= target_row:
            continue
        token.value = (
            f"{match.group('start')}:"
            f"{_formula_coordinate_with_row(match.group('end'), target_row)}"
        )
    return tokenizer.render()


def _shift_target_sheet_structures(
    ws,
    insert_row: int,
    amount: int,
    moved_merges: Sequence[CellRange],
) -> None:
    """Move structures which openpyxl's insert_rows deliberately leaves untouched."""
    for merged in moved_merges:
        ws.merge_cells(str(_shift_cell_range(merged, insert_row, amount)))

    moved_dimensions = {}
    for row_no, dimension in list(ws.row_dimensions.items()):
        if row_no >= insert_row:
            del ws.row_dimensions[row_no]
            moved = copy(dimension)
            moved.index = row_no + amount
            moved_dimensions[row_no + amount] = moved
    for row_no, dimension in moved_dimensions.items():
        ws.row_dimensions[row_no] = dimension

    if ws.auto_filter.ref:
        ws.auto_filter.ref = str(
            _shift_cell_range(CellRange(ws.auto_filter.ref), insert_row, amount, expand_adjacent=True)
        )
    if ws.print_area:
        shifted_areas = []
        for area in str(ws.print_area).split(","):
            local_area = area.rsplit("!", 1)[-1].replace("$", "")
            shifted_areas.append(
                str(_shift_cell_range(CellRange(local_area), insert_row, amount, expand_adjacent=True))
            )
        ws.print_area = ",".join(shifted_areas)
    for table in ws.tables.values():
        table.ref = str(_shift_cell_range(CellRange(table.ref), insert_row, amount, expand_adjacent=True))
    for validation in ws.data_validations.dataValidation:
        validation.sqref = _shift_multi_range(validation.sqref, insert_row, amount)

    conditional_rules = list(ws.conditional_formatting._cf_rules.items())
    ws.conditional_formatting._cf_rules.clear()
    for conditional, rules in conditional_rules:
        shifted = copy(conditional)
        shifted.sqref = _shift_multi_range(conditional.sqref, insert_row, amount)
        ws.conditional_formatting._cf_rules[shifted] = rules

    for row_break in ws.row_breaks.brk:
        if row_break.id >= insert_row:
            row_break.id += amount
    for drawing in [*ws._images, *ws._charts]:
        anchor = getattr(drawing, "anchor", None)
        for marker_name in ("_from", "_to"):
            marker = getattr(anchor, marker_name, None)
            if marker is not None and marker.row >= insert_row - 1:
                marker.row += amount


def append_records_to_workbook(
    template: str | Path | bytes | BinaryIO,
    output_path: str | Path | BinaryIO,
    records: Iterable[Mapping[str, Any]],
    field_aliases: Mapping[str, Sequence[str]],
    *,
    filename: str = "",
    sheet_names: Sequence[str] = (),
    field_formats: Mapping[str, str] | None = None,
    formula_fallback_fields: Collection[str] = (),
    record_overrides_formula_fields: Collection[str] = (),
    new_row_font_color: str = "",
    group_key_factory: Callable[[Mapping[str, Any]], str] | None = None,
    group_row_values_factory: Callable[[Mapping[str, Any], int], Mapping[int, Any]] | None = None,
    group_total_key_factory: Callable[[Mapping[str, Any]], str] | None = None,
    group_total_row_values_factory: Callable[
        [Sequence[Mapping[str, Any]], int, int, int, Mapping[int, str]],
        Mapping[int, Any],
    ] | None = None,
    first_total_section: bool = False,
) -> dict[str, Any]:
    """Copy the complete workbook and insert new detail rows before its total row.

    The supplied workbook is never modified in place.  Existing sheets, historical
    rows, formulas, drawings, page settings, defined names and workbook metadata are
    retained by editing the loaded workbook and saving it to ``output_path``.
    """
    workbook = load_complete_workbook_compatible(template, filename=filename)
    rows = [dict(record) for record in records]
    try:
        target, header_row, column_map = select_sheet(workbook, field_aliases, sheet_names)
        if not column_map:
            raise ValueError(f"{target.title} 未识别到任何可写入列")
        header_last_col = max(
            (
                col_no
                for col_no in range(1, (target.max_column or 1) + 1)
                if target.cell(header_row, col_no).value not in (None, "")
            ),
            default=max(column_map),
        )
        max_col = max(max(column_map), header_last_col)
        detail_columns = tuple(
            col_no
            for col_no, field in column_map.items()
            if field in _DETAIL_FIELDS
        ) or tuple(column_map)
        section_total_row: int | None = None
        reusable_blank_rows = 0
        if first_total_section:
            append_row, section_total_row, reusable_blank_rows = _first_total_section_rows(
                target,
                header_row,
                max_col,
                detail_columns=detail_columns,
            )
        else:
            append_row = _find_append_row(
                target,
                header_row,
                max_col,
                detail_columns=detail_columns,
            )
        style_row = _representative_data_row(target, header_row, append_row, max_col, column_map)
        inheritance = inherit_product_names(target, header_row, column_map, rows)
        grouped_layout: list[tuple[str, Any]] = []
        group_style_row: int | None = None
        group_total_style_row: int | None = None
        group_merges: list[CellRange] = []
        group_count = 0
        group_total_count = 0
        if group_total_key_factory or group_total_row_values_factory:
            if group_key_factory or group_row_values_factory:
                raise ValueError("分组标题行和分组合计行不能同时启用")
            if not group_total_key_factory or not group_total_row_values_factory:
                raise ValueError("分组合计行需要同时提供分组键和合计行写入规则")
            group_total_style_row = _nearest_total_row(
                target,
                header_row,
                append_row,
                max_col,
            )
            if group_total_style_row is None:
                group_total_style_row = style_row
            if append_row <= (target.max_row or append_row) and _row_has_total_marker(
                target,
                append_row,
                max_col,
            ):
                append_row += 1
            grouped_records: dict[str, list[dict[str, Any]]] = {}
            for index, record in enumerate(rows):
                key = str(group_total_key_factory(record) or "").strip() or f"__row_{index}"
                grouped_records.setdefault(key, []).append(record)
            for grouped_rows in grouped_records.values():
                grouped_layout.extend(("detail", record) for record in grouped_rows)
                grouped_layout.append(("total", grouped_rows))
            group_total_count = len(grouped_records)
        elif group_key_factory or group_row_values_factory:
            if not group_key_factory or not group_row_values_factory:
                raise ValueError("分组标题行需要同时提供分组键和标题行写入规则")
            group_style_row, group_merges = _nearest_group_title_row(
                target,
                header_row,
                style_row,
                max_col,
            )
            grouped_records: dict[str, list[dict[str, Any]]] = {}
            for index, record in enumerate(rows):
                key = str(group_key_factory(record) or "").strip() or f"__row_{index}"
                grouped_records.setdefault(key, []).append(record)
            for grouped_rows in grouped_records.values():
                grouped_layout.append(("group", grouped_rows[0]))
                grouped_layout.extend(("detail", record) for record in grouped_rows)
            group_count = len(grouped_records)
        else:
            grouped_layout = [("detail", record) for record in rows]

        formulas = []
        for sheet in workbook.worksheets:
            for cell in sheet._cells.values():
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append((sheet, cell.row, cell.column, cell.value))

        if grouped_layout:
            amount = len(grouped_layout)
            insert_row = section_total_row if section_total_row is not None else append_row
            insert_amount = (
                max(0, amount - reusable_blank_rows)
                if first_total_section
                else amount
            )
            moved_merges = [
                CellRange(str(value))
                for value in target.merged_cells.ranges
                if value.max_row >= insert_row
            ]
            if insert_amount:
                for merged in moved_merges:
                    target.unmerge_cells(str(merged))
                target.insert_rows(insert_row, insert_amount)
                _shift_target_sheet_structures(target, insert_row, insert_amount, moved_merges)

                for sheet, old_row, column, formula in formulas:
                    destination_row = old_row + insert_amount if sheet is target and old_row >= insert_row else old_row
                    sheet.cell(destination_row, column).value = _rewrite_formula_for_insert(
                        formula,
                        formula_sheet=sheet.title,
                        target_sheet=target.title,
                        formula_row=old_row,
                        insert_row=insert_row,
                        amount=insert_amount,
                    )

            pending_group_rows: list[int] = []
            for offset, (row_kind, payload) in enumerate(grouped_layout):
                row_no = append_row + offset
                source_row = (
                    group_style_row
                    if row_kind == "group"
                    else group_total_style_row
                    if row_kind == "total"
                    else style_row
                )
                if source_row is None:
                    raise ValueError("当前排期没有可复用的分组行")
                source_dimension = target.row_dimensions[source_row]
                if source_dimension.height:
                    target.row_dimensions[row_no].height = source_dimension.height
                for col_no in range(1, max_col + 1):
                    source_cell = target.cell(source_row, col_no)
                    destination = target.cell(row_no, col_no)
                    if isinstance(destination, MergedCell):
                        continue
                    _copy_style(source_cell, destination)
                    if isinstance(source_cell.value, str) and source_cell.value.startswith("="):
                        try:
                            destination.value = Translator(
                                source_cell.value,
                                origin=source_cell.coordinate,
                            ).translate_formula(destination.coordinate)
                        except (TypeError, ValueError):
                            destination.value = source_cell.value
                    if new_row_font_color:
                        font = copy(destination.font)
                        font.color = new_row_font_color
                        destination.font = font
                if row_kind == "group":
                    values = group_row_values_factory(payload, row_no)
                    if not isinstance(values, Mapping):
                        raise ValueError("分组标题行写入规则必须返回列号到值的映射")
                    for col_no, value in values.items():
                        target.cell(row_no, int(col_no)).value = _safe_value(value)
                    pending_group_rows.append(row_no)
                    continue
                if row_kind == "total":
                    grouped_rows = payload
                    group_end_row = row_no - 1
                    group_start_row = group_end_row - len(grouped_rows) + 1
                    values = group_total_row_values_factory(
                        grouped_rows,
                        group_start_row,
                        group_end_row,
                        row_no,
                        column_map,
                    )
                    if not isinstance(values, Mapping):
                        raise ValueError("分组合计行写入规则必须返回列号到值的映射")
                    for col_no, value in values.items():
                        target.cell(row_no, int(col_no)).value = _safe_value(value)
                    continue
                record = payload
                for col_no, field in column_map.items():
                    cell = target.cell(row_no, col_no)
                    if isinstance(cell, MergedCell):
                        continue
                    # A template formula is authoritative for calculated fields unless
                    # this customer explicitly identifies the field as PO-owned input.
                    has_formula = isinstance(cell.value, str) and cell.value.startswith("=")
                    if has_formula and field in record_overrides_formula_fields:
                        inherited_number_format = cell.number_format
                        safe_value = _safe_record_value(field, record.get(field, ""))
                        cell.value = safe_value
                        if isinstance(safe_value, datetime) and inherited_number_format == "General":
                            cell.number_format = "yyyy-mm-dd"
                    elif has_formula and field in formula_fallback_fields:
                        fallback = _safe_record_value(field, record.get(field, ""))
                        if fallback in (None, ""):
                            fallback_formula = '""'
                        elif isinstance(fallback, bool):
                            fallback_formula = "TRUE" if fallback else "FALSE"
                        elif isinstance(fallback, (int, float)):
                            fallback_formula = str(fallback)
                        else:
                            fallback_formula = '"' + str(fallback).replace('"', '""') + '"'
                        cell.value = f"=IFERROR({cell.value[1:]},{fallback_formula})"
                    elif not has_formula:
                        inherited_number_format = cell.number_format
                        safe_value = _safe_record_value(field, record.get(field, ""))
                        cell.value = safe_value
                        if isinstance(safe_value, datetime) and inherited_number_format == "General":
                            cell.number_format = "yyyy-mm-dd"
                    if field_formats and field in field_formats:
                        cell.number_format = field_formats[field]

            for row_no in pending_group_rows:
                for merged in group_merges:
                    target.merge_cells(
                        start_row=row_no,
                        end_row=row_no,
                        start_column=merged.min_col,
                        end_column=merged.max_col,
                    )

            first_total_row = (
                section_total_row + insert_amount
                if section_total_row is not None
                else append_row + amount
            )
            last_new_row = append_row + amount - 1
            for row_no in range(first_total_row, (target.max_row or first_total_row) + 1):
                for col_no in range(1, max_col + 1):
                    cell = target.cell(row_no, col_no)
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        cell.value = _extend_subtotal_to_previous_row(
                            cell.value,
                            last_new_row,
                        )

        if hasattr(output_path, "write"):
            workbook.save(output_path)
            output_reference = "<memory>"
        else:
            output_file = Path(output_path)
            output_file.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(output_file)
            output_reference = str(output_file)
        return {
            "path": output_reference,
            "template_sheet": target.title,
            "header_row": header_row,
            "style_row": style_row,
            "group_style_row": group_style_row,
            "group_total_style_row": group_total_style_row,
            "insert_row": append_row,
            "section_total_row": section_total_row,
            "reused_blank_rows": min(len(grouped_layout), reusable_blank_rows),
            "rows": len(rows),
            "group_rows": group_count,
            "group_total_rows": group_total_count,
            "inserted_rows": len(grouped_layout),
            "mapped_columns": {str(col): field for col, field in column_map.items()},
            "inheritance": inheritance,
            "mode": "full_workbook_append",
        }
    finally:
        workbook.close()


def append_column_records_to_workbook(
    template: str | Path | bytes | BinaryIO,
    output_path: str | Path | BinaryIO,
    records: Iterable[Any],
    *,
    filename: str = "",
    sheet_names: Sequence[str],
    header_row: int,
    max_col: int,
    detail_columns: Sequence[int] = (),
    row_values_factory: Callable[[Any, int], Mapping[int, Any]] | None = None,
    group_key_factory: Callable[[Any], str] | None = None,
    group_row_values_factory: Callable[[Any, int], Mapping[int, Any]] | None = None,
    column_formats: Mapping[int, str] | None = None,
    new_row_fill_color: str = "",
) -> dict[str, Any]:
    """Copy the complete workbook and append fixed-column records to one sheet.

    Every inserted row inherits the nearest preceding real detail row; total,
    subtotal, merged group-title, and blank separator rows are skipped.
    Values returned by ``row_values_factory`` may include row-relative formulas.
    Optional grouping inserts one copied item/title row before the matching detail
    rows. Empty values are intentionally skipped so copied template formulas and
    blanks remain authoritative.
    """
    workbook = load_complete_workbook_compatible(template, filename=filename)
    rows = list(records)
    try:
        target = next(
            (workbook[name] for name in sheet_names if name in workbook.sheetnames),
            None,
        )
        if target is None:
            raise ValueError(f"模板缺少目标工作表：{' / '.join(sheet_names)}")
        append_row = _find_append_row(
            target,
            header_row,
            max_col,
            detail_columns=detail_columns,
        )
        style_row = _nearest_detail_row(
            target,
            header_row,
            append_row,
            max_col,
        )
        grouped_layout: list[tuple[str, Any]] = []
        group_style_row: int | None = None
        group_merges: list[CellRange] = []
        group_count = 0
        if group_key_factory or group_row_values_factory:
            if not group_key_factory or not group_row_values_factory:
                raise ValueError("分组标题行需要同时提供分组键和标题行写入规则")
            group_style_row, group_merges = _nearest_group_title_row(
                target,
                header_row,
                style_row,
                max_col,
            )
            grouped_records: dict[str, list[Any]] = {}
            for index, record in enumerate(rows):
                key = str(group_key_factory(record) or "").strip() or f"__row_{index}"
                grouped_records.setdefault(key, []).append(record)
            for grouped_rows in grouped_records.values():
                grouped_layout.append(("group", grouped_rows[0]))
                grouped_layout.extend(("detail", record) for record in grouped_rows)
            group_count = len(grouped_records)
        else:
            grouped_layout = [("detail", record) for record in rows]
        highlight_fill = (
            PatternFill(fill_type="solid", fgColor=new_row_fill_color)
            if new_row_fill_color
            else None
        )

        formulas = []
        for sheet in workbook.worksheets:
            for cell in sheet._cells.values():
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append((sheet, cell.row, cell.column, cell.value))

        if grouped_layout:
            amount = len(grouped_layout)
            moved_merges = [
                CellRange(str(value))
                for value in target.merged_cells.ranges
                if value.max_row >= append_row
            ]
            for merged in moved_merges:
                target.unmerge_cells(str(merged))
            target.insert_rows(append_row, amount)
            _shift_target_sheet_structures(target, append_row, amount, moved_merges)

            for sheet, old_row, column, formula in formulas:
                destination_row = old_row + amount if sheet is target and old_row >= append_row else old_row
                sheet.cell(destination_row, column).value = _rewrite_formula_for_insert(
                    formula,
                    formula_sheet=sheet.title,
                    target_sheet=target.title,
                    formula_row=old_row,
                    insert_row=append_row,
                    amount=amount,
                )

            pending_group_rows: list[int] = []
            for offset, (row_kind, record) in enumerate(grouped_layout):
                row_no = append_row + offset
                source_row = group_style_row if row_kind == "group" else style_row
                if source_row is None:
                    raise ValueError("当前排期没有可复用的分组标题行")
                source_dimension = target.row_dimensions[source_row]
                if source_dimension.height:
                    target.row_dimensions[row_no].height = source_dimension.height
                for col_no in range(1, max_col + 1):
                    source_cell = target.cell(source_row, col_no)
                    destination = target.cell(row_no, col_no)
                    _copy_style(source_cell, destination)
                    if highlight_fill is not None:
                        destination.fill = copy(highlight_fill)
                    if isinstance(source_cell.value, str) and source_cell.value.startswith("="):
                        try:
                            destination.value = Translator(
                                source_cell.value,
                                origin=source_cell.coordinate,
                            ).translate_formula(destination.coordinate)
                        except (TypeError, ValueError):
                            destination.value = source_cell.value

                if row_kind == "group":
                    values = group_row_values_factory(record, row_no)
                    pending_group_rows.append(row_no)
                else:
                    values = row_values_factory(record, row_no) if row_values_factory else record
                if not isinstance(values, Mapping):
                    raise ValueError("固定列写入记录必须是列号到值的映射")
                for col_no, value in values.items():
                    if value in (None, ""):
                        continue
                    cell = target.cell(row_no, int(col_no))
                    cell.value = _safe_value(value)
                    if isinstance(value, (date, datetime)) and cell.number_format == "General":
                        cell.number_format = "yyyy-mm-dd"
                    if column_formats and int(col_no) in column_formats:
                        cell.number_format = column_formats[int(col_no)]

            for row_no in pending_group_rows:
                for merged in group_merges:
                    target.merge_cells(
                        start_row=row_no,
                        end_row=row_no,
                        start_column=merged.min_col,
                        end_column=merged.max_col,
                    )
                    if group_style_row is not None:
                        _copy_style(
                            target.cell(group_style_row, merged.min_col),
                            target.cell(row_no, merged.min_col),
                        )

            first_total_row = append_row + amount
            for row_no in range(first_total_row, (target.max_row or first_total_row) + 1):
                for col_no in range(1, max_col + 1):
                    cell = target.cell(row_no, col_no)
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        cell.value = _extend_subtotal_to_previous_row(
                            cell.value,
                            first_total_row - 1,
                        )

        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"

        if hasattr(output_path, "write"):
            workbook.save(output_path)
            output_reference = "<memory>"
        else:
            output_file = Path(output_path)
            output_file.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(output_file)
            output_reference = str(output_file)
        return {
            "path": output_reference,
            "template_sheet": target.title,
            "header_row": header_row,
            "style_row": style_row,
            "insert_row": append_row,
            "rows": len(rows),
            "group_rows": group_count,
            "inserted_rows": len(grouped_layout),
            "mode": "full_workbook_append",
        }
    finally:
        workbook.close()


def insert_column_records_after_matching_groups(
    template: str | Path | bytes | BinaryIO,
    output_path: str | Path | BinaryIO,
    records: Iterable[Any],
    *,
    filename: str = "",
    sheet_names: Sequence[str],
    header_row: int,
    max_col: int,
    existing_row_key_factory: Callable[[Any, int], str],
    record_key_factory: Callable[[Any], str],
    detail_columns: Sequence[int] = (),
    row_values_factory: Callable[[Any, int], Mapping[int, Any]] | None = None,
    column_formats: Mapping[int, str] | None = None,
) -> dict[str, Any]:
    """Insert fixed-column records after the last matching detail row.

    Records with the same non-empty key are inserted together immediately after
    the last existing row carrying that key, which keeps item-based schedules in
    ``detail... + subtotal`` order. Records without a matching key retain the
    established global append fallback. Formulas and workbook structures are
    shifted with Excel-like row insertion semantics.
    """
    workbook = load_complete_workbook_compatible(template, filename=filename)
    rows = list(records)
    try:
        target = next(
            (workbook[name] for name in sheet_names if name in workbook.sheetnames),
            None,
        )
        if target is None:
            raise ValueError(f"模板缺少目标工作表：{' / '.join(sheet_names)}")

        grouped_records: dict[str, list[Any]] = {}
        for index, record in enumerate(rows):
            key = str(record_key_factory(record) or "").strip()
            grouped_records.setdefault(key or f"__unmatched_{index}", []).append(record)

        last_matching_rows: dict[str, int] = {}
        for row_no in range(header_row + 1, (target.max_row or header_row) + 1):
            key = str(existing_row_key_factory(target, row_no) or "").strip()
            if key:
                last_matching_rows[key] = row_no

        global_append_row = _find_append_row(
            target,
            header_row,
            max_col,
            detail_columns=detail_columns,
        )
        fallback_style_row = _nearest_detail_row(
            target,
            header_row,
            global_append_row,
            max_col,
        )
        placements: dict[int, dict[str, Any]] = {}
        matched_groups = 0
        unmatched_groups = 0
        for key, grouped in grouped_records.items():
            matching_row = last_matching_rows.get(key)
            if matching_row is not None:
                insert_row = matching_row + 1
                style_row = matching_row
                matched_groups += 1
            else:
                insert_row = global_append_row
                style_row = fallback_style_row
                unmatched_groups += 1
            placement = placements.setdefault(
                insert_row,
                {"style_row": style_row, "records": []},
            )
            placement["records"].extend(grouped)

        inserted_rows = 0
        insertion_rows: list[int] = []
        for insert_row in sorted(placements, reverse=True):
            placement = placements[insert_row]
            placement_rows = placement["records"]
            amount = len(placement_rows)
            if not amount:
                continue
            style_row = int(placement["style_row"])
            formulas = [
                (sheet, cell.row, cell.column, cell.value)
                for sheet in workbook.worksheets
                for cell in sheet._cells.values()
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ]
            moved_merges = [
                CellRange(str(value))
                for value in target.merged_cells.ranges
                if value.max_row >= insert_row
            ]
            for merged in moved_merges:
                target.unmerge_cells(str(merged))
            target.insert_rows(insert_row, amount)
            _shift_target_sheet_structures(target, insert_row, amount, moved_merges)

            for sheet, old_row, column, formula in formulas:
                destination_row = (
                    old_row + amount
                    if sheet is target and old_row >= insert_row
                    else old_row
                )
                sheet.cell(destination_row, column).value = _rewrite_formula_for_insert(
                    formula,
                    formula_sheet=sheet.title,
                    target_sheet=target.title,
                    formula_row=old_row,
                    insert_row=insert_row,
                    amount=amount,
                )

            source_dimension = target.row_dimensions[style_row]
            for offset, record in enumerate(placement_rows):
                row_no = insert_row + offset
                if source_dimension.height:
                    target.row_dimensions[row_no].height = source_dimension.height
                for col_no in range(1, max_col + 1):
                    source_cell = target.cell(style_row, col_no)
                    destination = target.cell(row_no, col_no)
                    _copy_style(source_cell, destination)
                    if isinstance(source_cell.value, str) and source_cell.value.startswith("="):
                        try:
                            destination.value = Translator(
                                source_cell.value,
                                origin=source_cell.coordinate,
                            ).translate_formula(destination.coordinate)
                        except (TypeError, ValueError):
                            destination.value = source_cell.value

                values = row_values_factory(record, row_no) if row_values_factory else record
                if not isinstance(values, Mapping):
                    raise ValueError("固定列写入记录必须是列号到值的映射")
                for col_no, value in values.items():
                    if value in (None, ""):
                        continue
                    cell = target.cell(row_no, int(col_no))
                    cell.value = _safe_value(value)
                    if isinstance(value, (date, datetime)) and cell.number_format == "General":
                        cell.number_format = "yyyy-mm-dd"
                    if column_formats and int(col_no) in column_formats:
                        cell.number_format = column_formats[int(col_no)]

            inserted_rows += amount
            insertion_rows.append(insert_row)

        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"

        if hasattr(output_path, "write"):
            workbook.save(output_path)
            output_reference = "<memory>"
        else:
            output_file = Path(output_path)
            output_file.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(output_file)
            output_reference = str(output_file)
        return {
            "path": output_reference,
            "template_sheet": target.title,
            "header_row": header_row,
            "insert_rows": sorted(insertion_rows),
            "rows": len(rows),
            "matched_groups": matched_groups,
            "unmatched_groups": unmatched_groups,
            "inserted_rows": inserted_rows,
            "mode": "full_workbook_matching_group_insert",
        }
    finally:
        workbook.close()


def append_grouped_column_records_to_workbook(
    template: str | Path | bytes | BinaryIO,
    output_path: str | Path | BinaryIO,
    groups: Iterable[tuple[Any, Iterable[Any]]],
    *,
    filename: str = "",
    sheet_names: Sequence[str],
    header_row: int,
    max_col: int,
    summary_columns: Sequence[int],
    detail_style_columns: Sequence[int] = (),
    detail_row_values_factory: Callable[[Any, int], Mapping[int, Any]],
    summary_row_values_factory: Callable[[Any, list[Any], int, int, int], Mapping[int, Any]],
    reuse_trailing_summary_rows: bool = False,
    repair_existing_summary_values_factory: Callable[[int, int, int], Mapping[int, Any]] | None = None,
) -> dict[str, Any]:
    """Append grouped detail rows followed by one copied summary row per group.

    This is for schedules whose normal order pattern is ``detail... + item total``.
    The latest existing summary row supplies the summary style, while the nearest
    detail above it supplies the detail style. Templates with preformatted summary
    placeholders may reuse and clear those trailing rows instead of stacking new
    rows around them. Date number formats are deliberately left unchanged so the
    inserted rows display exactly like their template rows.
    """
    workbook = load_complete_workbook_compatible(template, filename=filename)
    grouped_rows = [(key, list(records)) for key, records in groups]
    grouped_rows = [(key, records) for key, records in grouped_rows if records]
    try:
        target = next(
            (workbook[name] for name in sheet_names if name in workbook.sheetnames),
            None,
        )
        if target is None:
            raise ValueError(f"模板缺少目标工作表：{' / '.join(sheet_names)}")

        def is_summary_row(row_no: int) -> bool:
            for col_no in summary_columns:
                value = target.cell(row_no, col_no).value
                if not isinstance(value, str) or value.startswith("="):
                    continue
                normalized = _normalise(value)
                if any(normalized == marker or normalized.endswith(marker) for marker in _TOTAL_MARKERS):
                    return True
            return False

        last_sheet_row = max(
            (
                cell.row
                for cell in target._cells.values()
                if cell.column <= max_col and cell.value not in (None, "")
            ),
            default=header_row,
        )
        trailing_summary_start: int | None = None
        if reuse_trailing_summary_rows:
            for row_no in range(last_sheet_row, header_row, -1):
                if not is_summary_row(row_no):
                    break
                trailing_summary_start = row_no
            if trailing_summary_start is None:
                raise ValueError(f"{target.title} 未找到可复用的尾部合计占位行")

        summary_style_search_end = (
            trailing_summary_start - 1
            if trailing_summary_start is not None
            else last_sheet_row
        )
        summary_style_row = next(
            (
                row_no
                for row_no in range(summary_style_search_end, header_row, -1)
                if is_summary_row(row_no)
            ),
            None,
        )
        if summary_style_row is None:
            raise ValueError(f"{target.title} 未找到可复用的分组合计行")
        detail_style_row = _nearest_detail_row(
            target,
            header_row,
            trailing_summary_start or summary_style_row,
            max_col,
        )
        if detail_style_columns:
            merged_rows = {
                row_no
                for merged in target.merged_cells.ranges
                if merged.max_col > merged.min_col
                for row_no in range(merged.min_row, merged.max_row + 1)
            }
            for candidate in range((trailing_summary_start or summary_style_row) - 1, header_row, -1):
                if candidate in merged_rows or is_summary_row(candidate):
                    continue
                values = [
                    target.cell(candidate, col_no).value
                    for col_no in range(1, max_col + 1)
                ]
                if not any(
                    value not in (None, "")
                    and not (isinstance(value, str) and value.startswith("="))
                    for value in values
                ):
                    continue
                style_cells = [
                    target.cell(candidate, int(col_no))
                    for col_no in detail_style_columns
                ]
                if any(
                    isinstance(cell.value, str) and cell.value.startswith("=")
                    for cell in style_cells
                ):
                    continue
                if any(cell.fill.fill_type not in (None, "none") for cell in style_cells):
                    continue
                detail_style_row = candidate
                break
        append_row = trailing_summary_start or (summary_style_row + 1)

        if repair_existing_summary_values_factory is not None:
            group_start_row = header_row + 1
            for row_no in range(header_row + 1, append_row):
                if not is_summary_row(row_no):
                    continue
                detail_rows = [
                    candidate
                    for candidate in range(group_start_row, row_no)
                    if not is_summary_row(candidate)
                    and any(
                        target.cell(candidate, col_no).value not in (None, "")
                        for col_no in range(1, max_col + 1)
                    )
                ]
                if detail_rows:
                    repaired_values = repair_existing_summary_values_factory(
                        row_no,
                        detail_rows[0],
                        detail_rows[-1],
                    )
                    if not isinstance(repaired_values, Mapping):
                        raise ValueError("既有合计修复规则必须返回列号到值的映射")
                    for col_no, value in repaired_values.items():
                        target.cell(row_no, int(col_no)).value = _safe_value(value)
                group_start_row = row_no + 1

        detail_template = [
            (
                target.cell(detail_style_row, col_no).coordinate,
                target.cell(detail_style_row, col_no).value,
                copy(target.cell(detail_style_row, col_no)._style),
            )
            for col_no in range(1, max_col + 1)
        ]
        summary_template = [
            (
                target.cell(summary_style_row, col_no).coordinate,
                target.cell(summary_style_row, col_no).value,
                copy(target.cell(summary_style_row, col_no)._style),
            )
            for col_no in range(1, max_col + 1)
        ]
        detail_row_height = target.row_dimensions[detail_style_row].height
        summary_row_height = target.row_dimensions[summary_style_row].height

        formulas = []
        for sheet in workbook.worksheets:
            for cell in sheet._cells.values():
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append((sheet, cell.row, cell.column, cell.value))

        amount = sum(len(records) + 1 for _key, records in grouped_rows)
        if amount:
            reusable_summary_rows = (
                last_sheet_row - append_row + 1
                if reuse_trailing_summary_rows
                else 0
            )
            insert_amount = (
                max(0, amount - reusable_summary_rows)
                if reuse_trailing_summary_rows
                else amount
            )
            insert_row = (
                last_sheet_row + 1
                if reuse_trailing_summary_rows
                else append_row
            )
            moved_merges = [
                CellRange(str(value))
                for value in target.merged_cells.ranges
                if value.max_row >= insert_row
            ]
            if insert_amount:
                for merged in moved_merges:
                    target.unmerge_cells(str(merged))
                target.insert_rows(insert_row, insert_amount)
                _shift_target_sheet_structures(target, insert_row, insert_amount, moved_merges)

                for sheet, old_row, column, formula in formulas:
                    destination_row = (
                        old_row + insert_amount
                        if sheet is target and old_row >= insert_row
                        else old_row
                    )
                    sheet.cell(destination_row, column).value = _rewrite_formula_for_insert(
                        formula,
                        formula_sheet=sheet.title,
                        target_sheet=target.title,
                        formula_row=old_row,
                        insert_row=insert_row,
                        amount=insert_amount,
                    )

            row_no = append_row
            for group_key, records in grouped_rows:
                detail_start_row = row_no
                for record in records:
                    target.row_dimensions[row_no].height = detail_row_height
                    for col_no, (origin, template_value, template_style) in enumerate(
                        detail_template,
                        start=1,
                    ):
                        destination = target.cell(row_no, col_no)
                        destination.value = None
                        destination._style = copy(template_style)
                        destination.comment = None
                        destination.hyperlink = None
                        if isinstance(template_value, str) and template_value.startswith("="):
                            try:
                                destination.value = Translator(
                                    template_value,
                                    origin=origin,
                                ).translate_formula(destination.coordinate)
                            except (TypeError, ValueError):
                                destination.value = template_value
                    values = detail_row_values_factory(record, row_no)
                    if not isinstance(values, Mapping):
                        raise ValueError("分组明细写入记录必须是列号到值的映射")
                    for col_no, value in values.items():
                        target.cell(row_no, int(col_no)).value = _safe_value(value)
                    row_no += 1

                detail_end_row = row_no - 1
                target.row_dimensions[row_no].height = summary_row_height
                for col_no, (origin, template_value, template_style) in enumerate(
                    summary_template,
                    start=1,
                ):
                    destination = target.cell(row_no, col_no)
                    destination.value = None
                    destination._style = copy(template_style)
                    destination.comment = None
                    destination.hyperlink = None
                    if isinstance(template_value, str) and template_value.startswith("="):
                        try:
                            destination.value = Translator(
                                template_value,
                                origin=origin,
                            ).translate_formula(destination.coordinate)
                        except (TypeError, ValueError):
                            destination.value = template_value
                values = summary_row_values_factory(
                    group_key,
                    records,
                    row_no,
                    detail_start_row,
                    detail_end_row,
                )
                if not isinstance(values, Mapping):
                    raise ValueError("分组合计写入记录必须是列号到值的映射")
                for col_no, value in values.items():
                    target.cell(row_no, int(col_no)).value = _safe_value(value)
                row_no += 1

            if reuse_trailing_summary_rows:
                clear_end_row = last_sheet_row + insert_amount
                for clear_row in range(append_row + amount, clear_end_row + 1):
                    target.row_dimensions[clear_row].height = None
                    for col_no in range(1, max_col + 1):
                        cell = target.cell(clear_row, col_no)
                        cell.value = None
                        cell._style = None
                        cell.comment = None
                        cell.hyperlink = None

            calculation = getattr(workbook, "calculation", None)
            if calculation is not None:
                calculation.fullCalcOnLoad = True
                calculation.forceFullCalc = True
                calculation.calcMode = "auto"

        if hasattr(output_path, "write"):
            workbook.save(output_path)
            output_reference = "<memory>"
        else:
            output_file = Path(output_path)
            output_file.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(output_file)
            output_reference = str(output_file)
        return {
            "path": output_reference,
            "template_sheet": target.title,
            "header_row": header_row,
            "detail_style_row": detail_style_row,
            "summary_style_row": summary_style_row,
            "insert_row": append_row,
            "detail_rows": sum(len(records) for _key, records in grouped_rows),
            "summary_rows": len(grouped_rows),
            "inserted_rows": amount,
            "reused_trailing_summary_rows": min(
                amount,
                last_sheet_row - append_row + 1,
            ) if reuse_trailing_summary_rows else 0,
            "mode": "full_workbook_grouped_append",
        }
    finally:
        workbook.close()


def create_new_order_workbook(
    template: str | Path | bytes | BinaryIO,
    output_path: str | Path | BinaryIO,
    records: Iterable[Mapping[str, Any]],
    field_aliases: Mapping[str, Sequence[str]],
    *,
    filename: str = "",
    sheet_names: Sequence[str] = (),
    header_rows: int | None = None,
    style_row: int | None = None,
    max_col: int | None = None,
    sheet_title: str = "新单",
    audit_sheet: bool = True,
    field_formats: Mapping[str, str] | None = None,
) -> dict[str, Any]:
    """Create a new-order-only workbook from the corresponding schedule template."""
    source_book = load_workbook_compatible(
        template,
        filename=filename,
        sheet_names=sheet_names,
    )
    rows = [dict(record) for record in records]
    try:
        source, detected_header, column_map = select_sheet(source_book, field_aliases, sheet_names)
        header_rows = header_rows or detected_header
        if not column_map:
            _, column_map = detect_header(source, field_aliases, scan_rows=max(20, header_rows))
        if not column_map:
            raise ValueError(f"{source.title} 未识别到任何可写入列")
        if max_col is None:
            header_last_col = max(
                (
                    col
                    for row in range(1, header_rows + 1)
                    for col in range(1, (source.max_column or 1) + 1)
                    if source.cell(row, col).value not in (None, "")
                ),
                default=max(column_map),
            )
            max_col = max(max(column_map), header_last_col)
        style_row = _style_source_row(source, header_rows, max_col, style_row)
        inheritance = inherit_product_names(
            source,
            header_rows,
            column_map,
            rows,
        )

        result = Workbook()
        output = result.active
        output.title = sheet_title[:31]
        _copy_headers(source, output, header_rows, max_col)

        for offset, record in enumerate(rows, start=header_rows + 1):
            if source.row_dimensions[style_row].height:
                output.row_dimensions[offset].height = source.row_dimensions[style_row].height
            for col in range(1, max_col + 1):
                old = source.cell(style_row, col)
                new = output.cell(offset, col)
                _copy_style(old, new)
                if isinstance(old.value, str) and old.value.startswith("="):
                    # Standalone exports intentionally omit historical/reference
                    # sheets.  Keep row-local formulas only; cross-sheet/external
                    # formulas would become broken links in the new workbook.
                    if "!" not in old.value and "[" not in old.value:
                        try:
                            new.value = Translator(old.value, origin=old.coordinate).translate_formula(new.coordinate)
                        except (TypeError, ValueError):
                            new.value = old.value
            for col, field in column_map.items():
                value = record.get(field, "")
                cell = output.cell(offset, col)
                cell.value = _safe_value(value)
                if field_formats and field in field_formats:
                    # 模板空白行有时残留了错误的日期/数字格式。字段语义应高于
                    # 样式行的 number_format，否则数量 800 会显示成 1902 年日期。
                    cell.number_format = field_formats[field]

        if audit_sheet:
            audit = result.create_sheet("识别明细")
            fields = list(field_aliases)
            for col, field in enumerate(fields, 1):
                cell = audit.cell(1, col, field)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for row_index, record in enumerate(rows, 2):
                for col, field in enumerate(fields, 1):
                    cell = audit.cell(row_index, col, _safe_value(record.get(field, "")))
                    if field_formats and field in field_formats:
                        cell.number_format = field_formats[field]
            audit.freeze_panes = "A2"
            for col, field in enumerate(fields, 1):
                width = max(
                    [len(str(field))]
                    + [len(str(rows[index].get(field, ""))) for index in range(min(len(rows), 100))]
                )
                audit.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 10), 42)

        if inheritance["applied"] or inheritance["conflicts"]:
            inherited = result.create_sheet("排期继承检查")
            headers = ("状态", "货号", "字段", "PO原值", "排期标准值 / 冲突值")
            for col, title in enumerate(headers, 1):
                cell = inherited.cell(1, col, title)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2F6B5F")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            row_no = 2
            for detail in inheritance["inherited_rows"]:
                values = (
                    "已按货号继承",
                    detail["item"],
                    detail["field"],
                    detail["po_value"],
                    detail["schedule_value"],
                )
                for col, value in enumerate(values, 1):
                    inherited.cell(row_no, col, _safe_value(value))
                row_no += 1
            for conflict in inheritance["conflicts"]:
                values = (
                    "排期存在冲突，未自动继承",
                    conflict["item"],
                    conflict["field"],
                    "",
                    " / ".join(conflict["values"]),
                )
                for col, value in enumerate(values, 1):
                    inherited.cell(row_no, col, _safe_value(value))
                row_no += 1
            inherited.freeze_panes = "A2"
            for col, width in enumerate((24, 20, 18, 36, 60), 1):
                inherited.column_dimensions[get_column_letter(col)].width = width

        if hasattr(output_path, "write"):
            result.save(output_path)
            output_reference = "<memory>"
        else:
            output_file = Path(output_path)
            output_file.parent.mkdir(parents=True, exist_ok=True)
            result.save(output_file)
            output_reference = str(output_file)
        result.close()
        return {
            "path": output_reference,
            "template_sheet": source.title,
            "header_rows": header_rows,
            "style_row": style_row,
            "rows": len(rows),
            "mapped_columns": {str(col): field for col, field in column_map.items()},
            "inheritance": inheritance,
            "mode": "new_order_only",
        }
    finally:
        source_book.close()
