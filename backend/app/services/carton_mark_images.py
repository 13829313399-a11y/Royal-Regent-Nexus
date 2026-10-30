"""Visible source images, thumbnails and local OCR with explicit reference provenance."""
import base64
import hashlib
import io
import json
import os
import subprocess
import sys
from collections import OrderedDict
import threading
import time
from pathlib import Path
from tempfile import TemporaryDirectory

from PIL import Image
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.services.carton_mark_ai_pilot import Evidence


def images(filename,content,worksheets=None):
    result=[];warnings=[]
    def append(data,location):
        if len(result)>=12:return
        try:
            with Image.open(io.BytesIO(data)) as picture:
                if picture.width*picture.height>40_000_000:raise ValueError("size")
                picture=picture.convert("RGB");picture.thumbnail((1400,1400))
                stream=io.BytesIO();picture.save(stream,format="JPEG",quality=85)
                thumb=picture.copy();thumb.thumbnail((360,260));preview=io.BytesIO();thumb.save(preview,format="JPEG",quality=80)
                result.append(dict(id=f"I{len(result)+1}",location=location,
                    data=base64.b64encode(stream.getvalue()).decode(),
                    preview="data:image/jpeg;base64,"+base64.b64encode(preview.getvalue()).decode()))
        except (OSError,ValueError):warnings.append(f"{location} 图片无法解码，需对照原稿核实。")
    if Path(filename).suffix.lower()==".xlsx":
        book=load_workbook(io.BytesIO(content),keep_links=False)
        try:
            for sheet in book:
                if sheet.sheet_state!="visible" or (worksheets and sheet.title not in worksheets):continue
                for index,picture in enumerate(sheet._images):
                    anchor=getattr(picture.anchor,"_from",None)
                    if anchor and (sheet.row_dimensions[anchor.row+1].hidden or sheet.column_dimensions[get_column_letter(anchor.col+1)].hidden):continue
                    append(picture._data(),f"{sheet.title}!image:{index+1}")
        finally:book.close()
    else:
        if worksheets:
            import xlrd
            book=xlrd.open_workbook(file_contents=content)
            visible={s.name for s in book.sheets() if not s.visibility}
            if set(worksheets)!=visible:return [],["旧版XLS按部分工作表读取时不读取其他工作表的图片，请对照所选原稿核实图示。"]
        # xlrd cannot expose legacy drawing pixels. A secure local Office export
        # supplies reference-page pixels; its OCR is still forbidden as business data.
        try:
            from app.services.document_tools.office_engine import render_office
            import pypdfium2
            with TemporaryDirectory(prefix="carton-images-") as temp:
                work=Path(temp);source=work/"source.xls";source.write_bytes(content);pdf=work/"source.pdf"
                render_office(source,pdf,{"secure_load":True,"recalculate":False},work,lambda:False)
                with pypdfium2.PdfDocument(pdf) as document:
                    for i in range(min(len(document),4)):
                        page=document[i];bitmap=page.render(scale=1.5);picture=bitmap.to_pil();stream=io.BytesIO()
                        try:picture.save(stream,format="PNG");append(stream.getvalue(),f"原稿打印参考!page:{i+1}")
                        finally:picture.close();bitmap.close();page.close()
        except Exception:warnings.append("旧版 XLS 的图片参考未完成本地读取，请对照原稿查看；原生单元格照常读取。")
    if len(result)==12:warnings.append("本次最多读取12张图片参考，其余图片请对照原稿。")
    return result,warnings


_CACHE=OrderedDict()
_CACHE_LOCK=threading.Lock()


def _ocr(serialized):
    key=hashlib.sha256(serialized).hexdigest()
    with _CACHE_LOCK:
        cached=_CACHE.get(key)
        if cached and time.monotonic()-cached[0]<3600:
            _CACHE.move_to_end(key);return cached[1]
    completed=subprocess.run([sys.executable,"-m","app.services.carton_mark_image_worker"],input=serialized,
        stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=60,
        cwd=str(Path(__file__).resolve().parents[2]),env={**os.environ,"PYTHONIOENCODING":"utf-8"},
        creationflags=getattr(subprocess,"CREATE_NO_WINDOW",0))
    output=completed.stdout.decode("utf-8",errors="replace")
    marker="CARTON_OCR_RESULT="
    if completed.returncode or marker not in output:raise ValueError("local OCR unavailable")
    result=json.loads(output.rsplit(marker,1)[1])
    with _CACHE_LOCK:
        _CACHE[key]=(time.monotonic(),result)
        while len(_CACHE)>64:_CACHE.popitem(last=False)
    return result


def collect(filename,content,worksheets=None,*,include_pixels=False):
    entries,warnings=images(filename,content,worksheets)
    if not entries:return [],[],warnings
    try:
        data=json.dumps([dict(id=e["id"],data=e["data"]) for e in entries]).encode()
        if len(data)>8_000_000:raise ValueError("size")
        readings={e["id"]:e["lines"] for e in _ocr(data)}
    except (ValueError,OSError,subprocess.TimeoutExpired):
        readings={};warnings.append("图片文字的本地识别未完成，已保留参考图供核对；未把图片数字作为本单字段。")
    notes=[];public=[]
    for entry in entries:
        lines=readings.get(entry["id"],[])
        text="\n".join(line["text"] for line in lines)[:2500]
        if text or include_pixels:
            notes.append(Evidence(id=entry["id"],kind="image_reference_text",location=entry["location"],
                text="图片参考，仅供解释制版要求，不能提供当前PO、货号、数量、重量或条码：\n"+(text or "本地文字未识别，请结合对应图片核对。")))
        public.append(dict(id=entry["id"],location=entry["location"],preview=entry["preview"],text=text,
                           status="needs_confirmation",sha256=hashlib.sha256(entry["data"].encode()).hexdigest()))
        if include_pixels:
            public[-1]['_pixels'] = entry['data']
    return notes,public,warnings
