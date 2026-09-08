from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


TEMPLATE_VERSION = "2026.07"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLS_CONTENT_TYPE = "application/vnd.ms-excel"
FIXED_TEMPLATE_DIRECTORY = (
    Path(__file__).resolve().parent.parent
    / "data"
    / "internal_quote_import_templates"
)
FIXED_TEMPLATE_FILE_NAMES = {
    "mold": "展兴模具--工模报价表.xlsx",
    "assembly": "装工.xlsx",
    "hardware": "五金1.xlsx",
    "painting": "喷油报价单.xlsx",
    "electronic": "电子报价单.xlsx",
    "sewing": "车缝报价单.xlsx",
    "hair": "车发部报价单.xls",
}

TEMPLATE_LABELS = {
    "mold": "工程部模具报价",
    "hardware": "工程部五金报价",
    "electronic": "电子部报价",
    "molding": "啤机部报价",
    "painting": "喷油部报价",
    "slush": "搪胶部报价",
    "sewing": "车缝部报价",
    "hair": "车发部报价",
    "assembly": "装配部排拉工序",
}

TEMPLATE_HEADERS = {
    "mold": [
        "序号", "模号", "模胚类型", "模具结构", "模具名称", "中文名称", "材质", "料型", "颜色",
        "出模数", "套数", "模具尺寸", "净重(g)", "模胚材质", "周期(秒)", "机型(TON)",
        "目标数", "模具规格", "工艺", "模价 RMB", "图片", "备注",
    ],
    "hardware": ["零件名称", "规格", "用量", "单价 RMB", "备注"],
    "electronic": ["零件名称", "规格", "用量", "单价 RMB", "金额 RMB", "税点 %", "备注"],
    "painting": [
        "图片引用", "名称", "位置",
        "夹模数量", "夹模单价 HKD", "移印数量", "移印单价 HKD",
        "UV数量", "UV单价 HKD",
        "散枪数量", "散枪单价 HKD", "边模数量", "边模单价 HKD",
        "油色数量", "油色单价 HKD", "浸油数量", "浸油单价 HKD",
        "抹油数量", "抹油单价 HKD", "擦PP水数量", "擦PP水单价 HKD", "备注",
    ],
    "slush": ["产品编号", "胶件名称", "材料", "料重(g)", "日产量24H", "用量(PC)", "单价 HKD", "备注"],
    "sewing": [
        "物料名称", "裁片部位", "供应商", "布料MOQ/Y", "低于MOQ/每色产生费用",
        "用量/码", "单价 RMB", "汇率", "成本 HKD（自动）", "码点", "价钱 HKD（自动）", "备注",
    ],
    "assembly": ["工序名称（组装）", "人数", "备注", "工序名称（包装/混装）", "人数", "备注"],
}

MOLDING_INJECTION_HEADERS = [
    "模具名称", "模号", "材质", "料型", "颜色", "啤净重(g)", "机台", "机型 A码",
    "出模数", "套数", "目标数", "周期(秒)", "成品用量", "备注", "料损耗 3%",
]
MOLDING_BLOW_HEADERS = [
    "货名", "日产量/22H", "用料", "料型", "预估料重 g", "吹工 HKD", "披锋 HKD",
    "利润×", "成品用量", "出数", "模价 RMB", "备注",
]

HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(color="FFFFFF", bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF7CC")
SECTION_FILL = PatternFill("solid", fgColor="CCFBF1")
THIN_BORDER = Border(
    left=Side(style="thin", color="D7E1E7"),
    right=Side(style="thin", color="D7E1E7"),
    top=Side(style="thin", color="D7E1E7"),
    bottom=Side(style="thin", color="D7E1E7"),
)


def _style_header(sheet, row_number: int, headers: list[str]) -> None:
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(row=row_number, column=column, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
        width = min(max(len(title) * 2 + 4, 13), 26)
        sheet.column_dimensions[get_column_letter(column)].width = max(
            sheet.column_dimensions[get_column_letter(column)].width or 0,
            width,
        )
    sheet.row_dimensions[row_number].height = 32


def _style_input_rows(sheet, start_row: int, end_row: int, column_count: int) -> None:
    for row in range(start_row, end_row + 1):
        for column in range(1, column_count + 1):
            cell = sheet.cell(row=row, column=column)
            cell.fill = INPUT_FILL
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[row].height = 24


def _add_instructions(workbook: Workbook, label: str) -> None:
    sheet = workbook.create_sheet("填写说明")
    rows = [
        [f"{label}导入模板", f"映射版本 {TEMPLATE_VERSION}"],
        ["填写规则", "在黄色区域填写数据；绿色表头名称不要修改、合并或删除。"],
        ["导入规则", "回到对应部门报价明细，点击“上传附件”，系统会自动识别并显示导入预览。"],
        ["金额规则", "模板中的自动金额仅供核对；正式金额保存后仍由服务端按冻结汇率与公式重算。"],
        ["空白字段", "无内容的非必填字段可以留空；请不要在数据区域插入合计或说明行。"],
    ]
    for row in rows:
        sheet.append(row)
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 88
    for row in range(1, len(rows) + 1):
        sheet.cell(row=row, column=1).font = Font(bold=True, color="0F766E")
        sheet.cell(row=row, column=1).fill = SECTION_FILL
        for column in (1, 2):
            sheet.cell(row=row, column=column).border = THIN_BORDER
            sheet.cell(row=row, column=column).alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"


def _build_standard_sheet(workbook: Workbook, import_type: str) -> None:
    headers = TEMPLATE_HEADERS[import_type]
    sheet = workbook.active
    sheet.title = TEMPLATE_LABELS[import_type]
    _style_header(sheet, 1, headers)
    _style_input_rows(sheet, 2, 21, len(headers))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}21"

    if import_type == "electronic":
        summary_column = len(headers) + 3
        value_column = summary_column + 1
        sheet.cell(row=1, column=summary_column, value="成本补充（RMB）")
        sheet.cell(row=1, column=summary_column).fill = HEADER_FILL
        sheet.cell(row=1, column=summary_column).font = HEADER_FONT
        sheet.cell(row=1, column=value_column, value="金额")
        sheet.cell(row=1, column=value_column).fill = HEADER_FILL
        sheet.cell(row=1, column=value_column).font = HEADER_FONT
        for row, label in enumerate(("邦定成本", "贴片成本", "人工成本", "测试费用", "包装运输", "抵税差额"), start=2):
            sheet.cell(row=row, column=summary_column, value=label)
            sheet.cell(row=row, column=value_column).fill = INPUT_FILL
            sheet.cell(row=row, column=summary_column).border = THIN_BORDER
            sheet.cell(row=row, column=value_column).border = THIN_BORDER
        sheet.cell(row=9, column=summary_column, value="10%利润（可修改百分比）")
        sheet.column_dimensions[get_column_letter(summary_column)].width = 24
        sheet.column_dimensions[get_column_letter(value_column)].width = 14


def _build_molding_sheet(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = TEMPLATE_LABELS["molding"]
    _style_header(sheet, 1, MOLDING_INJECTION_HEADERS)
    _style_input_rows(sheet, 2, 16, len(MOLDING_INJECTION_HEADERS))
    blow_header_row = 18
    sheet.cell(row=17, column=1, value="吹气部分（没有吹气报价可留空）")
    sheet.cell(row=17, column=1).font = Font(bold=True, color="0F766E")
    sheet.cell(row=17, column=1).fill = SECTION_FILL
    _style_header(sheet, blow_header_row, MOLDING_BLOW_HEADERS)
    _style_input_rows(sheet, blow_header_row + 1, blow_header_row + 15, len(MOLDING_BLOW_HEADERS))
    sheet.freeze_panes = "A2"


def build_internal_quote_import_template(import_type: str) -> tuple[bytes, str]:
    label = TEMPLATE_LABELS.get(import_type)
    if label is None:
        raise ValueError("不支持的内部报价导入模板类型")

    fixed_file_name = FIXED_TEMPLATE_FILE_NAMES.get(import_type)
    if fixed_file_name is not None:
        template_path = FIXED_TEMPLATE_DIRECTORY / fixed_file_name
        try:
            content = template_path.read_bytes()
        except OSError as error:
            raise RuntimeError(
                f"内部报价固定模板缺失：{fixed_file_name}"
            ) from error
        signature = b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1" if template_path.suffix.lower() == ".xls" else b"PK"
        if not content.startswith(signature):
            raise RuntimeError(
                f"内部报价固定模板格式无效：{fixed_file_name}"
            )
        return content, fixed_file_name

    workbook = Workbook()
    if import_type == "molding":
        _build_molding_sheet(workbook)
    else:
        _build_standard_sheet(workbook, import_type)
    _add_instructions(workbook, label)
    workbook.active = 0

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue(), f"{label}导入模板-{TEMPLATE_VERSION}.xlsx"
