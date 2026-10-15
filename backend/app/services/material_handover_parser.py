"""Layout-independent field mapping over native/OCR tables, never posting stock."""
import io
import math
import re
import tempfile
import unicodedata
from pathlib import Path
from decimal import Decimal, InvalidOperation
from zipfile import ZipFile, BadZipFile

from app.schemas.material_handover import HandoverData, HandoverLine
from app.services.document_tools.document_ir import DocumentIR, Cell, Table, SourceAnchor, Block, ToolError

MAX_BYTES = 20 * 1024 * 1024
SUFFIXES = {'.xlsx', '.xls', '.docx', '.pdf', '.png', '.jpg', '.jpeg', '.webp'}
ALIASES = {
    'item_code': ('物料编码', '材料编码', '物料代码', '材料代码', '物料编号', '料号', '物料料号', 'materialcode', 'itemcode'),
    'item_name': ('用料名称', '物料名称', '材料名称', '品名', '布料名称色号', '布料名称', 'materialname', 'description'),
    'quantity': ('本次送货数量', '送货数量', '发料数量', '实发数量', '领料数量', '合计数量', '总数量', '数量', '计划用量', '采购数量', 'qty', 'quantity'),
    'unit': ('领料单位', '单位', '采购单位', 'unit', 'uom'),
    'specification': ('规格', '尺寸', '规格型号', 'specification'),
    'color': ('颜色', '色号', 'color', 'colour'),
    'lot': ('缸号', '批号', '批次号', '批次', 'lot'),
    'roll_no': ('卷号', '卷支编号', 'rollno'),
    'contract_no': ('合同号', '合同编号', 'contractno'),
    'production_no': ('放产订单', '放产单号', '生产单号', '生产清单号', '生产单编号', 'productionno'),
    'product_no': ('货号', '款号', '产品编码', 'productcode'),
    'note': ('备注', '说明', 'remark'),
    'unit_price': ('单价', '含税单价', '不含税单价', 'unitprice', 'price'),
    'price_unit': ('计价单位', '价格单位', 'priceunit'),
    'currency': ('币种', '货币', 'currency'),
    'price_note': ('单价依据', '价格说明', 'pricenote'),
}
UNIT_ALIASES = {'y': '码', 'yd': '码', 'yds': '码', 'yard': '码', 'yards': '码', 'kg': '千克', 'kgs': '千克', 'pcs': 'Pcs', 'pc': 'Pcs', 'm': '米'}


def normalized(value):
    return re.sub(r'[^A-Za-z0-9\u4e00-\u9fff]', '', unicodedata.normalize('NFKC', str(value or ''))).lower()


def numeric(value):
    match = re.fullmatch(r'\s*([0-9]+(?:,[0-9]{3})*(?:\.[0-9]{1,6})?)\s*([A-Za-z\u4e00-\u9fff]*)\s*', str(value))
    if not match:
        return '', ''
    try:
        amount = Decimal(match[1].replace(',', ''))
        if not amount.is_finite() or amount <= 0:
            return '', ''
        return format(amount, 'f'), UNIT_ALIASES.get(match[2].lower(), match[2])
    except InvalidOperation:
        return '', ''


def header_label(cell, column):
    if not cell:
        return ''
    parts = [normalized(v) for v in cell.raw_text.splitlines() if v.strip()]
    aliases = {a for values in ALIASES.values() for a in values}
    # OCR can merge two neighbouring header cells while their data rows remain separate.
    if cell.colspan > 1 and len(parts) == cell.colspan and (all(v.isdigit() for v in parts) or any(v in aliases for v in parts)):
        return parts[column - cell.column]
    return normalized(cell.raw_text)


def validate_file(name, content):
    name = Path(name.replace('\\', '/')).name
    if not name or len(name) > 255 or Path(name).suffix.lower() not in SUFFIXES:
        raise ToolError('FILE_TYPE', '支持 Excel、Word DOCX、PDF、JPG、PNG 和 WebP，请核对文件格式')
    if not 0 < len(content) <= MAX_BYTES:
        raise ToolError('FILE_SIZE', '单个文档须为 1 字节至 20 MB')
    if Path(name).suffix.lower() in {'.xlsx', '.docx'}:
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(i.file_size for i in entries) > 100 * 1024 * 1024:
                    raise ToolError('FILE_SIZE', '文档解压后过大，请拆分后上传')
                if any('vbaproject' in i.filename.lower() for i in entries):
                    raise ToolError('FILE_TYPE', '不接受带宏的文档，请另存为普通 Excel / DOCX')
        except BadZipFile as exc:
            raise ToolError('FILE_TYPE', 'Office 文件结构无效，请重新导出') from exc
    return name


def image_ir(content):
    from PIL import Image, ImageOps
    from app.services.document_tools import local_ocr
    with Image.open(io.BytesIO(content)) as original:
        if original.width * original.height > 40_000_000 or getattr(original, 'n_frames', 1) != 1:
            raise ToolError('IMAGE_SIZE', '请使用不超过 4000 万像素的静态图片')
        image = ImageOps.exif_transpose(original).convert('RGB')
        ratio = min(1, math.sqrt(6_000_000 / max(1, image.width * image.height)))
        if ratio < 1:
            image = image.resize((int(image.width * ratio), int(image.height * ratio)))
        box = [0, 0, image.width, image.height]
        words = local_ocr.recognize(image, box)
        tables = local_ocr.recover_tables(image, words, box, 0, 'photo')
        ir = DocumentIR(source_type='image', tables=tables)
        ir.blocks = [Block(id=f'text-{i}', text=w['text'], source=SourceAnchor(page_index=0, bbox_pt=w['bbox_pt'], method='local_ocr')) for i, w in enumerate(words)]
        image.close()
        return ir


def extract_ir(name, content, folder):
    name = validate_file(name, content)
    path = folder / ('source' + Path(name).suffix.lower())
    path.write_bytes(content)
    suffix = path.suffix
    if suffix in {'.png', '.jpg', '.jpeg', '.webp'}:
        return image_ir(content)
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted or not 1 <= len(reader.pages) <= 20:
            raise ToolError('PDF_PAGES', '请使用未加密且不超过 20 页的 PDF')
        from app.services.document_tools.pdf_engine import _extract
        # Uploaded handover papers use local OCR only; no automatic paid request.
        return _extract(path, {'ai_mode': 'off'}, folder, lambda *args: None, lambda: False, ocr=True).ir
    if suffix == '.xls':
        import xlrd
        book = xlrd.open_workbook(file_contents=content)
        ir = DocumentIR(source_type='xls')
        for sheet in book.sheets():
            if sheet.nrows * sheet.ncols > 50000:
                raise ToolError('TABLE_SIZE', '工作表太大，请拆分后上传')
            table = Table(id=sheet.name, row_count=sheet.nrows, column_count=sheet.ncols, source=SourceAnchor(sheet=sheet.name))
            for r in range(sheet.nrows):
                for c in range(sheet.ncols):
                    value = sheet.cell_value(r, c)
                    text = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
                    table.cells.append(Cell(id=f'{r}-{c}', row=r, column=c, raw_text=text, value=text, source=SourceAnchor(sheet=sheet.name, cell=f'R{r+1}C{c+1}')))
            ir.tables.append(table)
        return ir
    from app.services.document_tools.office_engine import _extract_xlsx, _extract_docx
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        with path.open('rb') as stream:
            book = load_workbook(stream, read_only=True, keep_links=False)
            if sum(s.max_row * s.max_column for s in book) > 50000:
                raise ToolError('TABLE_SIZE', '工作表太大，请拆分后上传')
            book.close()
        return _extract_xlsx(path, {}, lambda *args: None, lambda: False)
    return _extract_docx(path, {}, lambda *args: None, lambda: False)


def map_ir(ir):
    lines, evidence, warnings = [], {}, []
    texts = [b.text for b in ir.blocks if b.text]
    texts += [c.raw_text for t in ir.tables for c in t.cells if c.raw_text]
    full_text = '\n'.join(texts)
    budget = any(v in full_text for v in ('预算用料', '计划用量', '放产损耗'))
    data = HandoverData(kind='ISSUE' if budget or any(v in full_text for v in ('发料单', '领料单', '领料申请')) else 'DELIVERY')
    if budget:
        warnings.append('这是用料计划 / 预算候选，请核对实际申请量；生成交接单不会直接收发库存。')
    for key, aliases in {'source_no': ('单号', '单据编号', '生产清单号', '生产单号', 'D No'), 'contract_no': ('合同号',), 'product_no': ('款号', '货号')}.items():
        for label in aliases:
            pattern = r'D\s*N[oe0]' if label == 'D No' else re.escape(label)
            match = re.search(pattern + r'\s*[:：.]?\s*([A-Za-z0-9][A-Za-z0-9_\-./]{2,127})', full_text, re.I)
            if match:
                setattr(data, key, match[1]); break
    date_match = re.search(r'(20\d{2})[年/\-.](\d{1,2})[月/\-.](\d{1,2})', full_text)
    if date_match:
        from datetime import date
        try:
            data.document_date = date(*map(int, date_match.groups()))
        except ValueError:
            warnings.append('单据日期无法确认，请核对原件。')
    for table in ir.tables:
        if table.row_count > 1000 or table.column_count > 100:
            warnings.append('有表格超出识别范围，请拆分或补录。'); continue
        grid = {}
        for cell in table.cells:
            for r in range(cell.row, cell.row + cell.rowspan):
                for c in range(cell.column, cell.column + cell.colspan):
                    grid[r, c] = cell
        mapping, start = {}, None
        for r in range(min(25, table.row_count)):
            candidate = {}
            for c in range(table.column_count):
                cell = grid.get((r, c))
                label = header_label(cell, c)
                for field, aliases in ALIASES.items():
                    if label in aliases:
                        rank = aliases.index(label)
                        if field not in candidate or rank < candidate[field][1]:
                            candidate[field] = (c, rank)
            extra_header = False
            if 'quantity' not in candidate and ('item_code' in candidate or 'item_name' in candidate) and r + 1 < table.row_count:
                for c in range(table.column_count):
                    below = grid.get((r+1, c))
                    above = grid.get((r, c))
                    label = header_label(below, c)
                    combined = normalized((above.raw_text if above else '') + (below.raw_text if below else ''))
                    for value in (combined, label):
                        if value in ALIASES['quantity']:
                            rank = ALIASES['quantity'].index(value)
                            if 'quantity' not in candidate or rank < candidate['quantity'][1]:
                                candidate['quantity'] = (c, rank); extra_header = True
            if 'quantity' in candidate and ('item_code' in candidate or 'item_name' in candidate):
                # Repeated 单位 columns belong to different objects; prefer the material name's nearest unit.
                units = [(c, ALIASES['unit'].index(normalized(grid[r,c].raw_text))) for c in range(table.column_count)
                    if (r,c) in grid and normalized(grid[r,c].raw_text) in ALIASES['unit']]
                if units:
                    material_column = candidate.get('item_name', candidate.get('item_code'))[0]
                    candidate['unit'] = min(units, key=lambda pair: (pair[1], abs(pair[0] - material_column)))
                mapping, start = candidate, r + 1 + int(extra_header)
                # Subheaders preserve outer columns such as 合计数量 / 单位.
                if not extra_header and r + 1 < table.row_count:
                    next_cells = [header_label(grid.get((r+1,c)), c) for c in range(table.column_count)]
                    identity_values = [(next_cells[candidate[f][0]], ALIASES[f]) for f in ('item_code', 'item_name') if f in candidate]
                    if sum(v.isdigit() for v in next_cells) >= 2 and not any(v and v not in aliases for v, aliases in identity_values):
                        start += 1
                break
        if start is None:
            continue
        for r in range(start, table.row_count):
            raw, anchors = {}, {}
            for field, (c, _) in mapping.items():
                cell = grid.get((r, c))
                raw[field] = (cell.display_text or cell.raw_text).strip() if cell and cell.resolution not in {'unknown', 'unresolved'} and (field != 'quantity' or cell.row == r) else ''
                if cell:
                    anchors[field] = {'raw': cell.raw_text[:2000], 'display': cell.display_text[:2000], 'source': cell.source.model_dump(), 'resolution': cell.resolution}
            if not raw.get('item_code') and not raw.get('item_name'):
                continue
            if normalized(raw.get('item_name')) in {'合计', '总计', '小计', '物料名称', '材料名称'}:
                continue
            if not raw.get('item_code'):
                found = re.search(r'(?<!\d)\d{2}(?:-\d{2}){1,3}-\d{4,6}(?!\d)', raw.get('item_name', ''))
                if found:
                    raw['item_code'] = found[0]
            if not raw.get('item_name') and not raw.get('item_code'):
                continue
            amount, suffix_unit = numeric(raw.get('quantity', ''))
            raw['quantity'] = amount
            if 'unit_price' in raw:
                price=raw['unit_price'].replace(',','').strip()
                if not re.fullmatch(r'[0-9]+(?:\.[0-9]{1,6})?',price):
                    if price:warnings.append(f'第 {len(lines)+1} 行单价不清楚，请对照原件补充。')
                    price=''
                raw['unit_price']=price
            if 'currency' in raw:
                currency={'人民币':'CNY','港币':'HKD','美元':'USD','欧元':'EUR'}.get(raw['currency'],raw['currency'].upper())
                raw['currency']=currency if currency in {'CNY','HKD','USD','EUR'} else ''
            raw['unit'] = UNIT_ALIASES.get(raw.get('unit', '').lower(), raw.get('unit', '')) or suffix_unit
            if suffix_unit and raw['unit'] != suffix_unit:
                warnings.append(f'第 {len(lines)+1} 行数量内单位与单位列不一致，请核对。')
                raw['unit'] = ''
            identifier = f'L{len(lines)+1}'
            for key, value in raw.items():
                limit = next((getattr(m, 'max_length', None) for m in HandoverLine.model_fields[key].metadata if getattr(m, 'max_length', None)), 2000)
                if len(value) > limit:
                    warnings.append(f'第 {len(lines)+1} 行字段较长，请对照原件核对完整说明。')
                    raw[key] = value[:limit]
            row = HandoverLine(id=identifier, source_line=f'{table.title or table.id}:{r+1}'[:128], **raw)
            numeric_rolls = []
            used_roll_cells = set()
            quantity_cell = grid.get((r, mapping['quantity'][0]))
            for c in range(table.column_count):
                header = grid.get((start-1, c))
                value = grid.get((r, c))
                # OCR may span a total cell into the neighbouring numbered-roll column.
                # Count each physical roll cell once, and never reuse the row total.
                if value is quantity_cell or c in {column for column, _ in mapping.values()}:
                    continue
                if header and header_label(header, c).isdigit() and value and value.id not in used_roll_cells and value.resolution not in {'unknown', 'unresolved'} and numeric(value.raw_text)[0]:
                    numeric_rolls.append(value.raw_text.strip())
                    used_roll_cells.add(value.id)
            row.roll_quantities = ' / '.join(numeric_rolls)
            row_text = '\n'.join(grid[r,c].raw_text for c in range(table.column_count) if (r,c) in grid)
            if not row.lot:
                lot_match = re.search(r'缸号\s*[:：]?\s*([A-Za-z0-9#_\-]+)', row_text)
                if lot_match:
                    row.lot = lot_match[1]
            if numeric_rolls and amount and sum(Decimal(numeric(v)[0]) for v in numeric_rolls) != Decimal(amount):
                warnings.append(f'第 {len(lines)+1} 行逐卷数量合计与单据总数不同，请核对原件和未辨认的卷数。')
            if not row.item_code or not row.item_name or not row.quantity or not row.unit:
                warnings.append(f'第 {len(lines)+1} 行缺少清楚的物料编码、名称、数量或单位，请补充核对。')
            lines.append(row); evidence[identifier] = anchors
            if len(lines) > 200:
                raise ToolError('TOO_MANY_LINES', '单据超过 200 行，请拆分后上传')
    data.lines = lines
    warnings += [issue.message for issue in ir.issues][:30]
    if ir.source_type in {'image', 'pdf'}:
        warnings.append('请核对印刷文字与手写 / 涂改内容，数量和单位以确认后的明细为准。')
    if not lines:
        warnings.append('未识别到完整物料表，请查看原件并补录，或重新上传更清晰的文件。')
    return {'data': data.model_dump(mode='json'), 'warnings': list(dict.fromkeys(warnings))[:60], 'evidence': evidence, 'raw_text': full_text[:100000], 'engine': ir.source_type}


def parse(name, content):
    with tempfile.TemporaryDirectory(prefix='rr-material-extract-') as folder:
        return map_ir(extract_ir(name, content, Path(folder)))
