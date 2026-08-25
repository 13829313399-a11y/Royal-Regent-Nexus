from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
from pathlib import PurePosixPath
import math
import re
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from lxml import etree
import msoffcrypto
import openpyxl
import xlrd

from app.services.customer_order_manual import (
    apply_overrides_to_preview,
    decorate_manual_resolution_policy,
)


XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PREVIEW_SCHEMA_VERSION = "customer-order-buzzbee-preview-v1"
STANDARD_TEMPLATE = "BUZZBEE_STANDARD_CONTRACT_V1"
WMC_TEMPLATE = "BUZZBEE_WALMART_WMC_INLINE_V1"
WMU_TEMPLATE = "BUZZBEE_WALMART_WMU_PO_ATTACHED_V1"
TARGET_TEMPLATE = "BUZZBEE_PRODUCTION_SCHEDULE_V1"
SCHEDULE_PASSWORD = "2026"
MAX_PO_BYTES = 12 * 1024 * 1024
MAX_SCHEDULE_BYTES = 35 * 1024 * 1024
MAX_BATCH_PO_FILES = 30
MAX_BATCH_PO_BYTES = 80 * 1024 * 1024
ITEM_SHEET_BULLET = "子弹枪ITEM表"
ITEM_SHEET_WATER = "水枪ITEM表"
ITEM_BOUNDARY_KEYWORDS = ("取消单", "转单", "已走货")
SKIPPABLE_BLOCKER_LABELS = {
    "missing_product_name_zh": "中文名称留空，稍后由跟客补充",
    "missing_unit_price": "单价及金额留空，稍后由跟客补充",
    "existing_order_line": "测试阶段确认重复导入当前排期已有订单",
}

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"m": MAIN_NS, "r": OFFICE_REL_NS, "p": PACKAGE_REL_NS}
CELL_REF_PATTERN = re.compile(r"^([A-Z]{1,3})(\d+)$")
FORMULA_CELL_REF_PATTERN = re.compile(r"(?<![A-Z0-9_])(\$?[A-Z]{1,3})(\$?)(\d+)")
FORMULA_RANGE_PATTERN = re.compile(
    r"(?<![A-Z0-9_])(\$?[A-Z]{1,3})(\$?)(\d+):(\$?[A-Z]{1,3})(\$?)(\d+)"
)
PACKAGING_PATTERN = re.compile(r"\b(\d{5}-\d{2}-\d{2}-[A-Za-z0-9]+)\b")
PO_PATTERN = re.compile(r"(?:P\s*/?\s*O|PO)\s*(?:NUMBER|NO\.?|/[^:]+)?\s*:\s*([0-9]{7,})", re.I)

LABEL_FIELDS = (
    ("Contract No.", "contract_no"),
    ("Date of Loading", "requested_ship_date"),
    ("Final Inspection Date", "inspection_date"),
    ("Inspection Date", "inspection_date"),
    ("Our Item#", "product_no"),
    ("Goods", "product_name_en"),
    ("Quantity", "quantity"),
    ("Export Carton Packing", "units_per_carton"),
    ("Shipping Carton Packing", "units_per_carton"),
)


class CustomerOrderWorkbookError(ValueError):
    pass


@dataclass
class ParsedPoLine:
    values: dict[str, Any]
    lineage: dict[str, str]
    input_template: str
    source_file_name: str = ""


@dataclass
class ScheduleLookup:
    product_name_zh: str = ""
    unit_price_hkd: Decimal | None = None
    existing_rows: list[dict[str, str]] | None = None


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _clean_identifier(value: Any) -> str:
    text = _clean_text(value)
    if not text:
        return ""
    try:
        numeric = Decimal(text.replace(",", ""))
    except InvalidOperation:
        return text
    if numeric == numeric.to_integral_value():
        return str(numeric.quantize(Decimal("1")))
    return text


def _decimal(value: Any) -> Decimal | None:
    text = _clean_text(value).replace(",", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _decimal_text(value: Decimal | None, places: int | None = None) -> str:
    if value is None:
        return ""
    normalized = value if places is None else value.quantize(Decimal(1).scaleb(-places))
    text = format(normalized, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _format_iso_date(value: Any, *, datemode: int = 0) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float, Decimal)) and value > 10000:
        try:
            return xlrd.xldate_as_datetime(float(value), datemode).date().isoformat()
        except (ValueError, OverflowError):
            return ""

    text = _clean_text(value)
    if not text:
        return ""
    upper = text.upper().replace(",", " ").replace(".", " ")
    if "TO BE ADVISED" in upper or "TBA" == upper.strip():
        return ""

    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%d-%b-%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(upper, pattern).date().isoformat()
        except ValueError:
            pass

    match = re.search(
        r"\b(\d{1,2})\s+(JAN(?:UARY)?|FEB(?:RUARY)?|MAR(?:CH)?|APR(?:IL)?|MAY|"
        r"JUN(?:E)?|JUL(?:Y)?|AUG(?:UST)?|SEP(?:T(?:EMBER)?)?|OCT(?:OBER)?|"
        r"NOV(?:EMBER)?|DEC(?:EMBER)?)\s+(\d{2,4})\b",
        upper,
    )
    if match:
        year = int(match.group(3))
        if year < 100:
            year += 2000
        month = datetime.strptime(match.group(2)[:3], "%b").month
        return date(year, month, int(match.group(1))).isoformat()
    return ""


def _excel_serial(value: str) -> str:
    parsed = date.fromisoformat(value)
    return str((parsed - date(1899, 12, 30)).days)


def _label_matches(text: str, label: str) -> bool:
    normalized = text.strip()
    return (
        normalized == label
        or normalized.startswith(f"{label} ")
        or normalized.startswith(f"{label}:")
        or normalized.startswith(f"{label} :")
    )


def _next_non_empty(values: list[tuple[int, Any]], start_col: int) -> tuple[int, Any] | None:
    for column, value in values:
        if column <= start_col:
            continue
        if _clean_text(value):
            return column, value
    return None


def _cell_name(row: int, column: int) -> str:
    return f"{openpyxl.utils.get_column_letter(column)}{row}"


def _ordinary_customer_metadata(
    file_name: str,
    rows: list[tuple[int, list[tuple[int, Any, str]]]],
    *,
    contract_row: int | None,
    contract_value_column: int | None,
) -> tuple[str, str, str, str, str]:
    row_values = {
        row_number: {
            column: (_clean_text(value), coordinate)
            for column, value, coordinate in row
            if _clean_text(value)
        }
        for row_number, row in rows
    }
    customer_fragments: list[str] = []
    customer_sources: list[str] = []
    if contract_row is not None and contract_value_column is not None:
        candidates = [
            (column, text, coordinate)
            for column, (text, coordinate) in row_values.get(contract_row, {}).items()
            if contract_value_column < column <= contract_value_column + 6
            and re.search(r"[A-Za-z]", text)
            and not any(_label_matches(text, label) for label, _ in LABEL_FIELDS)
        ]
        if candidates:
            customer_column, current_fragment, current_source = candidates[-1]
            previous = row_values.get(contract_row - 1, {}).get(customer_column)
            if previous and re.search(r"[A-Za-z]", previous[0]) and not re.search(
                r"\b(?:CONTACT|DATE|PURCHASE ORDER)\b", previous[0], re.I
            ):
                customer_fragments.append(previous[0])
                customer_sources.append(previous[1])
            customer_fragments.append(current_fragment)
            customer_sources.append(current_source)

    raw_customer = re.sub(r"\s+", " ", " ".join(customer_fragments)).strip(" -/")
    all_text = " ".join(
        _clean_text(value)
        for _, row in rows
        for _, value, _ in row
        if _clean_text(value)
    )
    probe = re.sub(r"\s+", " ", f"{file_name} {raw_customer} {all_text}").upper()

    customer_name = raw_customer
    country = ""
    profile_standard = ""
    profile_source = ""
    if "DOLLAR GENERAL" in probe:
        customer_name, country, profile_standard = "Dollar General", "美国", "美国标准"
        profile_source = "客户规则 · Dollar General"
    elif re.search(r"\bA{2,3}FE(?:S)?\b", probe):
        customer_name, country, profile_standard = "AAFES", "美国", "美国标准"
        profile_source = "客户规则 · AAFES"
    elif "SAFARI HOUSE" in probe:
        customer_name, country, profile_standard = "SAFARI HOUSE", "科威特", "欧洲标准"
        profile_source = "客户规则 · SAFARI HOUSE"
    elif "TOTTUS" in probe:
        if "CHILE" in probe:
            customer_name, country = "TOTTUS -CHILE", "智利"
        elif "PERU" in probe:
            customer_name, country = "TOTTUS -PERU", "秘鲁"
        else:
            customer_name = "TOTTUS"
        profile_standard = "欧洲标准"
        profile_source = "客户规则 · TOTTUS"
    elif "ABU ISSA" in probe:
        customer_name, country, profile_standard = "ABU ISSA", "卡塔尔", "欧洲标准"
        profile_source = "客户规则 · ABU ISSA"
    elif re.search(r"\bSKME\b", probe):
        customer_name, profile_standard = "SKME", "欧洲标准"
        profile_source = "客户规则 · SKME"
    elif re.search(r"\bCOOP\b", probe):
        customer_name, profile_standard = "COOP", "欧洲标准"
        profile_source = "客户规则 · COOP"

    country_aliases = (
        ("UNITED STATES", "美国"),
        ("U.S.A.", "美国"),
        ("USA", "美国"),
        ("CANADA", "加拿大"),
        ("KUWAIT", "科威特"),
        ("QATAR", "卡塔尔"),
        ("PERU", "秘鲁"),
        ("CHILE", "智利"),
    )
    if not country:
        for token, localized in country_aliases:
            if re.search(rf"\b{re.escape(token)}\b", raw_customer, re.I):
                country = localized
                customer_name = re.sub(
                    rf"(?:\s*[-/]?\s*)\b{re.escape(token)}\b\s*$",
                    "",
                    customer_name,
                    flags=re.I,
                ).strip(" -/")
                break

    if "EUROPEAN STANDARD" in probe or "EUROPEAN SAFETY" in probe or "2009/48/EC" in probe:
        standard = "欧洲标准"
        standard_source = "合同条款 · European Standard"
    elif "CANADIAN STANDARD" in probe:
        standard = "加拿大标准"
        standard_source = "合同条款 · Canadian Standard"
    elif any(token in probe for token in ("AMERICAN STANDARD", "U.S. STANDARD", "ASTM F963")):
        standard = "美国标准"
        standard_source = "合同条款 · U.S. Standard"
    else:
        standard = profile_standard
        standard_source = profile_source

    customer_source = (
        " / ".join(customer_sources)
        if customer_sources
        else profile_source or "合同表头 · 未识别"
    )
    return customer_name, country, standard, customer_source, standard_source


def _parse_ordinary_po_rows(
    file_name: str,
    sheet_name: str,
    rows: list[tuple[int, list[tuple[int, Any, str]]]],
    *,
    datemode: int = 0,
) -> list[ParsedPoLine]:
    values: dict[str, Any] = {}
    lineage: dict[str, str] = {}
    field_locations: dict[str, tuple[int, int]] = {}
    packaging = ""
    packaging_source = ""
    po_no = ""
    po_source = ""

    for row_number, row in rows:
        simple_row = [(column, value) for column, value, _ in row]
        for column, raw_value, coordinate in row:
            text = _clean_text(raw_value)
            if not text:
                continue
            for label, field in LABEL_FIELDS:
                if field in values or not _label_matches(text, label):
                    continue
                found = _next_non_empty(simple_row, column)
                if found:
                    value_column, value = found
                    values[field] = value
                    lineage[field] = f"{sheet_name}!{_cell_name(row_number, value_column)}"
                    field_locations[field] = (row_number, value_column)
            if not packaging:
                match = PACKAGING_PATTERN.search(text)
                if match:
                    packaging = match.group(1)
                    packaging_source = f"{sheet_name}!{coordinate}"
            if not po_no and "PO NUMBER" in text.upper():
                found = _next_non_empty(simple_row, column)
                if found:
                    value_column, value = found
                    po_no = _clean_text(value)
                    po_source = f"{sheet_name}!{_cell_name(row_number, value_column)}"
            if not po_no:
                match = PO_PATTERN.search(text)
                if match:
                    po_no = match.group(1)
                    po_source = f"{sheet_name}!{coordinate}"

    contract_location = field_locations.get("contract_no")
    customer_name, country, standard, customer_source, standard_source = (
        _ordinary_customer_metadata(
            file_name,
            rows,
            contract_row=contract_location[0] if contract_location else None,
            contract_value_column=contract_location[1] if contract_location else None,
        )
    )
    if customer_source and "!" not in customer_source and re.fullmatch(
        r"[A-Z]+\d+(?: / [A-Z]+\d+)*", customer_source
    ):
        customer_source = " / ".join(
            f"{sheet_name}!{coordinate}" for coordinate in customer_source.split(" / ")
        )
    inspection_raw = _clean_text(values.get("inspection_date"))
    values["contract_no"] = _clean_identifier(values.get("contract_no"))
    values["product_no"] = _clean_identifier(values.get("product_no"))
    values["po_no"] = po_no
    values["customer_name"] = customer_name
    values["country"] = country
    values["standard"] = standard
    values["packaging"] = packaging
    values["requested_ship_date"] = _format_iso_date(
        values.get("requested_ship_date"), datemode=datemode
    )
    values["inspection_raw"] = inspection_raw
    values["inspection_date"] = _format_iso_date(
        values.get("inspection_date"), datemode=datemode
    )
    lineage.update(
        {
            "po_no": po_source,
            "customer_name": customer_source,
            "country": customer_source,
            "standard": standard_source,
            "packaging": packaging_source,
        }
    )
    return [ParsedPoLine(values=values, lineage=lineage, input_template=STANDARD_TEMPLATE)]


def _parse_xls_po(file_name: str, content: bytes) -> list[ParsedPoLine]:
    try:
        workbook = xlrd.open_workbook(file_contents=content, formatting_info=True)
    except Exception as exc:
        raise CustomerOrderWorkbookError(f"普通 PO 文件无法读取：{exc}") from exc
    sheet = workbook.sheet_by_index(0)
    rows = [
        (
            row_index + 1,
            [
                (
                    column + 1,
                    sheet.cell_value(row_index, column),
                    _cell_name(row_index + 1, column + 1),
                )
                for column in range(sheet.ncols)
            ],
        )
        for row_index in range(sheet.nrows)
    ]
    return _parse_ordinary_po_rows(
        file_name,
        sheet.name,
        rows,
        datemode=workbook.datemode,
    )


def _iter_xlsx_rows(sheet: openpyxl.worksheet.worksheet.Worksheet):
    for row_number, row in enumerate(sheet.iter_rows(), start=1):
        yield row_number, [
            (column, cell.value, _cell_name(row_number, column))
            for column, cell in enumerate(row, start=1)
        ]


def _parse_wmc_xlsx_sheet(
    sheet: openpyxl.worksheet.worksheet.Worksheet,
) -> list[ParsedPoLine]:
    values: dict[str, Any] = {}
    lineage: dict[str, str] = {}
    packaging = ""
    packaging_source = ""
    po_no = ""
    po_source = ""
    inspection_raw = ""
    inspection_source = ""

    for row_number, row in _iter_xlsx_rows(sheet):
        simple_row = [(column, value) for column, value, _ in row]
        for column, raw_value, coordinate in row:
            text = _clean_text(raw_value)
            if not text:
                continue
            for label, field in LABEL_FIELDS:
                if field in values or not _label_matches(text, label):
                    continue
                found = _next_non_empty(simple_row, column)
                if found:
                    value_column, value = found
                    values[field] = value
                    lineage[field] = f"{sheet.title}!{_cell_name(row_number, value_column)}"
            if not packaging:
                match = PACKAGING_PATTERN.search(text)
                if match:
                    packaging = match.group(1)
                    packaging_source = f"{sheet.title}!{coordinate}"
            if not po_no:
                match = PO_PATTERN.search(text)
                if match:
                    po_no = match.group(1)
                    po_source = f"{sheet.title}!{coordinate}"
            if not inspection_raw and "INSPECTION" in text.upper():
                inspection_raw = text.split(":", 1)[-1].strip()
                inspection_source = f"{sheet.title}!{coordinate}"

    values["contract_no"] = _clean_identifier(values.get("contract_no") or sheet["N4"].value)
    values["product_no"] = _clean_identifier(values.get("product_no") or sheet["F12"].value)
    values["product_name_en"] = _clean_text(values.get("product_name_en") or sheet["C14"].value)
    values["quantity"] = values.get("quantity") or sheet["F16"].value
    values["po_no"] = po_no
    values["customer_name"] = "WMC"
    values["country"] = "加拿大"
    values["standard"] = "加拿大标准"
    values["packaging"] = packaging
    values["requested_ship_date"] = _format_iso_date(values.get("requested_ship_date") or sheet["F8"].value)
    values["inspection_raw"] = inspection_raw
    values["inspection_date"] = _format_iso_date(inspection_raw)
    lineage.update(
        {
            "contract_no": lineage.get("contract_no") or f"{sheet.title}!N4",
            "product_no": lineage.get("product_no") or f"{sheet.title}!F12",
            "product_name_en": lineage.get("product_name_en") or f"{sheet.title}!C14",
            "quantity": lineage.get("quantity") or f"{sheet.title}!F16",
            "po_no": po_source,
            "customer_name": "模板规则 · P4=WMC",
            "country": "模板规则 · WMC → 加拿大",
            "standard": "模板规则 · WMC → 加拿大标准",
            "packaging": packaging_source,
            "inspection_date": inspection_source,
        }
    )
    return [ParsedPoLine(values=values, lineage=lineage, input_template=WMC_TEMPLATE)]


def _parse_wmu_xlsx_workbook(
    workbook: openpyxl.Workbook,
) -> list[ParsedPoLine]:
    """Expand a mainland Walmart WMU workbook from its PO Attached sheet."""
    attached = next(
        (sheet for sheet in workbook.worksheets if sheet.title.strip().lower() == "po attached"),
        None,
    )
    if attached is None:
        raise CustomerOrderWorkbookError("WMU 合同缺少 PO Attached 子订单页")

    main = workbook.worksheets[0]
    product_no = _clean_identifier(main["E14"].value or attached["B2"].value)
    product_name_en = _clean_text(main["E17"].value or attached["B3"].value)
    units_per_carton = _decimal(main["E25"].value)
    packaging_refs = [
        _clean_text(main.cell(row=row_number, column=1).value)
        for row_number in range(1, main.max_row + 1)
    ]
    packaging = " / ".join(
        value for value in packaging_refs if re.search(r"\d{4,}-\d{2}-\d{2}-WM", value, re.I)
    )

    header_row = next(
        (
            row_number
            for row_number in range(1, min(attached.max_row, 20) + 1)
            if "WALMART PO" in _clean_text(attached.cell(row_number, 2).value).upper()
            and "ORDERED QTY" in _clean_text(attached.cell(row_number, 6).value).upper()
        ),
        None,
    )
    if header_row is None:
        raise CustomerOrderWorkbookError("WMU PO Attached 页未识别到子订单表头")

    lines: list[ParsedPoLine] = []
    for row_number in range(header_row + 1, attached.max_row + 1):
        contract_no = _clean_identifier(attached.cell(row_number, 1).value)
        po_no = _clean_text(attached.cell(row_number, 2).value)
        quantity = _decimal(attached.cell(row_number, 6).value)
        carton_count = _decimal(attached.cell(row_number, 7).value)
        if not contract_no and not po_no:
            continue
        if not contract_no or not po_no or quantity is None or quantity <= 0:
            continue
        row_pack = units_per_carton
        if row_pack is None and carton_count is not None and carton_count > 0:
            row_pack = quantity / carton_count
        values = {
            "contract_no": contract_no,
            "po_no": po_no,
            "customer_name": "WALMART USA",
            "country": "美国",
            "product_no": product_no,
            "product_name_en": product_name_en,
            "quantity": quantity,
            "units_per_carton": row_pack,
            "standard": "美国标准",
            "packaging": packaging or "Walmart USA / RFID",
            "requested_ship_date": _format_iso_date(attached.cell(row_number, 4).value),
            "inspection_raw": _clean_text(attached.cell(row_number, 8).value),
            "inspection_date": _format_iso_date(attached.cell(row_number, 8).value),
            "ship_via": _clean_text(attached.cell(row_number, 3).value),
        }
        lineage = {
            "contract_no": f"{attached.title}!A{row_number}",
            "po_no": f"{attached.title}!B{row_number}",
            "customer_name": f"{attached.title}!A1",
            "country": "模板规则 · WALMART USA → 美国",
            "product_no": f"{main.title}!E14",
            "product_name_en": f"{main.title}!E17",
            "quantity": f"{attached.title}!F{row_number}",
            "units_per_carton": f"{main.title}!E25",
            "standard": "模板规则 · WALMART USA → 美国标准",
            "packaging": f"{main.title}!A52:A53" if packaging else "模板规则 · WMU RFID",
            "requested_ship_date": f"{attached.title}!D{row_number}",
            "inspection_date": f"{attached.title}!H{row_number}",
        }
        lines.append(ParsedPoLine(values=values, lineage=lineage, input_template=WMU_TEMPLATE))

    if not lines:
        raise CustomerOrderWorkbookError("WMU PO Attached 页没有可用的子订单明细")
    return lines


def _parse_xlsx_po(file_name: str, content: bytes) -> list[ParsedPoLine]:
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise CustomerOrderWorkbookError(f"PO 文件无法读取：{exc}") from exc
    try:
        sheet = workbook.worksheets[0]
        if _clean_text(sheet["P4"].value).upper() == "WMC":
            return _parse_wmc_xlsx_sheet(sheet)

        rows = list(_iter_xlsx_rows(sheet))
        identity_probe = " ".join(
            [file_name]
            + [
                _clean_text(value)
                for _, row in rows
                for _, value, _ in row
                if _clean_text(value)
            ]
        ).upper()
        if "WMU" in identity_probe and any(
            item.title.strip().lower() == "po attached" for item in workbook.worksheets
        ):
            return _parse_wmu_xlsx_workbook(workbook)
        if "INDONESIA" in identity_probe or "印尼" in identity_probe:
            raise CustomerOrderWorkbookError("印尼合同不属于当前 BuzzBee 映射范围")
        if "WMU" in identity_probe:
            raise CustomerOrderWorkbookError("WMU 合同缺少 PO Attached 子订单页")
        return _parse_ordinary_po_rows(file_name, sheet.title, rows)
    finally:
        workbook.close()


def parse_po(file_name: str, content: bytes) -> list[ParsedPoLine]:
    lowered = file_name.lower()
    if lowered.endswith(".xls") and not lowered.endswith(".xlsx"):
        lines = _parse_xls_po(file_name, content)
    elif lowered.endswith((".xlsx", ".xlsm")):
        lines = _parse_xlsx_po(file_name, content)
    else:
        raise CustomerOrderWorkbookError("PO 文件只支持 .xls 或 .xlsx")
    for line in lines:
        line.source_file_name = file_name
    return lines


def _decrypt_schedule(content: bytes) -> tuple[bytes, bool]:
    try:
        office = msoffcrypto.OfficeFile(BytesIO(content))
    except Exception as exc:
        raise CustomerOrderWorkbookError("客户排期不是有效的 Excel 工作簿") from exc
    if not office.is_encrypted():
        return content, False
    try:
        office.load_key(password=SCHEDULE_PASSWORD)
        output = BytesIO()
        office.decrypt(output)
        return output.getvalue(), True
    except Exception as exc:
        raise CustomerOrderWorkbookError("客户排期无法用已配置的 BuzzBee 密码打开") from exc


def _encrypt_schedule(content: bytes) -> bytes:
    try:
        office = msoffcrypto.OfficeFile(BytesIO(content))
        output = BytesIO()
        office.encrypt(SCHEDULE_PASSWORD, output)
        return output.getvalue()
    except Exception as exc:
        raise CustomerOrderWorkbookError(f"生成的客户排期重新加密失败：{exc}") from exc


class OoxmlSchedule:
    def __init__(self, content: bytes):
        try:
            with ZipFile(BytesIO(content)) as archive:
                self.parts = {name: archive.read(name) for name in archive.namelist()}
                self.part_info = {name: archive.getinfo(name) for name in archive.namelist()}
        except Exception as exc:
            raise CustomerOrderWorkbookError("解密后的客户排期不是有效的 XLSX 文件") from exc
        self.shared_strings = self._read_shared_strings()
        self.sheet_paths = self._read_sheet_paths()
        self._validate_template()

    def _read_shared_strings(self) -> list[str]:
        payload = self.parts.get("xl/sharedStrings.xml")
        if not payload:
            return []
        root = etree.fromstring(payload)
        return [
            "".join(node.text or "" for node in string_item.iter(f"{{{MAIN_NS}}}t"))
            for string_item in root.findall("m:si", NS)
        ]

    def _read_sheet_paths(self) -> dict[str, str]:
        workbook = etree.fromstring(self.parts["xl/workbook.xml"])
        relationships = etree.fromstring(self.parts["xl/_rels/workbook.xml.rels"])
        targets = {
            item.get("Id"): item.get("Target")
            for item in relationships.findall("p:Relationship", NS)
        }
        result: dict[str, str] = {}
        for sheet in workbook.findall("m:sheets/m:sheet", NS):
            target = targets.get(sheet.get(f"{{{OFFICE_REL_NS}}}id"))
            if not target:
                continue
            posix_target = PurePosixPath(target.replace("\\", "/"))
            result[sheet.get("name", "")] = str(
                posix_target if str(posix_target).startswith("xl/") else PurePosixPath("xl") / posix_target
            ).lstrip("/")
        return result

    def _validate_template(self) -> None:
        required = {"接单表", "正单评审表", ITEM_SHEET_BULLET, ITEM_SHEET_WATER}
        missing = required - self.sheet_paths.keys()
        if missing:
            raise CustomerOrderWorkbookError(
                "排期结构不匹配 BUZZBEE_PRODUCTION_SCHEDULE_V1，"
                f"缺少工作表：{', '.join(sorted(missing))}"
            )
        order_rows = self.read_rows("接单表", limit=3)
        review_rows = self.read_rows("正单评审表", limit=3)
        bullet_item_rows = self.read_rows(ITEM_SHEET_BULLET, limit=3)
        water_item_rows = self.read_rows(ITEM_SHEET_WATER, limit=3)
        if (
            order_rows.get(3, {}).get("C") != "P/O#:"
            or review_rows.get(3, {}).get("C") != "PO.NO"
            or "PO.NO" not in bullet_item_rows.get(3, {}).get("C", "")
            or "PO.NO" not in water_item_rows.get(3, {}).get("C", "")
        ):
            raise CustomerOrderWorkbookError(
                "排期表头与当前 BuzzBee 模板不一致；印尼排期或其他版本不能用于本次输出"
            )

    def _cell_value(self, cell: etree._Element) -> str:
        cell_type = cell.get("t")
        if cell_type == "inlineStr":
            return "".join(node.text or "" for node in cell.iter(f"{{{MAIN_NS}}}t"))
        value = cell.find("m:v", NS)
        if value is None:
            return ""
        raw = value.text or ""
        if cell_type == "s":
            try:
                return self.shared_strings[int(raw)]
            except (ValueError, IndexError):
                return raw
        return raw

    def read_rows(self, sheet_name: str, *, limit: int | None = None) -> dict[int, dict[str, str]]:
        root = etree.fromstring(self.parts[self.sheet_paths[sheet_name]])
        rows: dict[int, dict[str, str]] = {}
        for row in root.findall("m:sheetData/m:row", NS):
            row_number = int(row.get("r", "0"))
            if limit is not None and row_number > limit:
                continue
            values: dict[str, str] = {}
            for cell in row.findall("m:c", NS):
                match = CELL_REF_PATTERN.match(cell.get("r", ""))
                if match:
                    values[match.group(1)] = self._cell_value(cell)
            rows[row_number] = values
        return rows

    def build_schedule_index(self) -> dict[str, ScheduleLookup]:
        index: dict[str, ScheduleLookup] = {}
        for row_number, values in self.read_rows("接单表").items():
            if row_number <= 3:
                continue
            if "合计" in values.get("M", ""):
                break
            product_no = _clean_identifier(values.get("F"))
            if not product_no:
                continue
            record = index.setdefault(product_no, ScheduleLookup(existing_rows=[]))
            if values.get("G"):
                record.product_name_zh = values["G"].strip()
            price = _decimal(values.get("M"))
            if price is not None and price > 0:
                record.unit_price_hkd = price
            record.existing_rows.append(
                {
                    "row": str(row_number),
                    "po_no": values.get("C", "").strip(),
                    "contract_no": _clean_identifier(values.get("D")),
                    "product_no": product_no,
                    "requested_ship_date": _format_iso_date(_decimal(values.get("S")) or values.get("S")),
                }
            )
        return index

    def _find_marker_row(self, sheet_name: str, column: str, marker: str) -> int:
        for row_number, values in self.read_rows(sheet_name).items():
            if marker in values.get(column, ""):
                return row_number
        raise CustomerOrderWorkbookError(f"{sheet_name} 未找到“{marker}”插入边界")

    @staticmethod
    def _shift_cell_reference(reference: str, insert_row: int) -> str:
        match = CELL_REF_PATTERN.match(reference)
        if not match:
            return reference
        row_number = int(match.group(2))
        if row_number >= insert_row:
            row_number += 1
        return f"{match.group(1)}{row_number}"

    @staticmethod
    def _shift_formula(formula: str, insert_row: int) -> str:
        if not formula:
            return formula
        placeholders: list[str] = []

        def replace_range(match: re.Match[str]) -> str:
            start_column, start_abs, start_text, end_column, end_abs, end_text = match.groups()
            start_row = int(start_text)
            end_row = int(end_text)
            if start_row >= insert_row:
                start_row += 1
            if end_row >= insert_row:
                end_row += 1
            elif end_row == insert_row - 1:
                end_row += 1
            token = (
                f"{start_column}{start_abs}{start_row}:"
                f"{end_column}{end_abs}{end_row}"
            )
            placeholders.append(token)
            return f"__RANGE_{len(placeholders) - 1}__"

        shifted = FORMULA_RANGE_PATTERN.sub(replace_range, formula)

        def replace_cell(match: re.Match[str]) -> str:
            column, absolute_marker, row_text = match.groups()
            row_number = int(row_text)
            if row_number >= insert_row:
                row_number += 1
            return f"{column}{absolute_marker}{row_number}"

        shifted = FORMULA_CELL_REF_PATTERN.sub(replace_cell, shifted)
        for index, token in enumerate(placeholders):
            shifted = shifted.replace(f"__RANGE_{index}__", token)
        return shifted

    @staticmethod
    def _shift_range_reference(reference: str, insert_row: int) -> str:
        tokens = reference.split()
        shifted_tokens: list[str] = []
        for token in tokens:
            if ":" in token:
                start, end = token.split(":", 1)
                shifted_start = OoxmlSchedule._shift_cell_reference(start.replace("$", ""), insert_row)
                shifted_end = OoxmlSchedule._shift_cell_reference(end.replace("$", ""), insert_row)
                if CELL_REF_PATTERN.match(start.replace("$", "")) and int(
                    CELL_REF_PATTERN.match(end.replace("$", "")).group(2)
                ) == insert_row - 1:
                    end_match = CELL_REF_PATTERN.match(shifted_end)
                    shifted_end = f"{end_match.group(1)}{int(end_match.group(2)) + 1}"
                shifted_tokens.append(f"{shifted_start}:{shifted_end}")
            else:
                shifted_tokens.append(OoxmlSchedule._shift_cell_reference(token.replace("$", ""), insert_row))
        return " ".join(shifted_tokens)

    @staticmethod
    def _clear_cell(cell: etree._Element) -> None:
        for child in list(cell):
            cell.remove(child)
        cell.attrib.pop("t", None)

    @staticmethod
    def _set_cell(
        row: etree._Element,
        column: str,
        row_number: int,
        *,
        value: str | int | Decimal | None = None,
        formula: str = "",
        cached: Decimal | int | None = None,
        inline: bool = False,
    ) -> None:
        reference = f"{column}{row_number}"
        cell = next((item for item in row.findall("m:c", NS) if item.get("r") == reference), None)
        if cell is None:
            cell = etree.Element(f"{{{MAIN_NS}}}c", r=reference)
            row.append(cell)
            row[:] = sorted(
                row,
                key=lambda item: openpyxl.utils.column_index_from_string(
                    CELL_REF_PATTERN.match(item.get("r", "A1")).group(1)
                )
                if item.tag == f"{{{MAIN_NS}}}c"
                else 100000,
            )
        OoxmlSchedule._clear_cell(cell)
        if formula:
            formula_node = etree.SubElement(cell, f"{{{MAIN_NS}}}f")
            formula_node.text = formula
            if cached is not None:
                value_node = etree.SubElement(cell, f"{{{MAIN_NS}}}v")
                value_node.text = _decimal_text(Decimal(cached))
            return
        if value is None or value == "":
            return
        if inline:
            cell.set("t", "inlineStr")
            inline_node = etree.SubElement(cell, f"{{{MAIN_NS}}}is")
            text_node = etree.SubElement(inline_node, f"{{{MAIN_NS}}}t")
            text_node.text = str(value)
            return
        value_node = etree.SubElement(cell, f"{{{MAIN_NS}}}v")
        value_node.text = _decimal_text(value if isinstance(value, Decimal) else Decimal(str(value)))

    def read_cell_formula(self, sheet_name: str, row_number: int, column: str) -> str:
        root = etree.fromstring(self.parts[self.sheet_paths[sheet_name]])
        row = next(
            (
                item
                for item in root.findall("m:sheetData/m:row", NS)
                if int(item.get("r", "0")) == row_number
            ),
            None,
        )
        if row is None:
            return ""
        reference = f"{column}{row_number}"
        cell = next((item for item in row.findall("m:c", NS) if item.get("r") == reference), None)
        if cell is None:
            return ""
        formula_node = cell.find("m:f", NS)
        return formula_node.text or "" if formula_node is not None else ""

    def set_cell(
        self,
        sheet_name: str,
        row_number: int,
        column: str,
        **payload: Any,
    ) -> None:
        path = self.sheet_paths[sheet_name]
        root = etree.fromstring(self.parts[path])
        row = next(
            (
                item
                for item in root.findall("m:sheetData/m:row", NS)
                if int(item.get("r", "0")) == row_number
            ),
            None,
        )
        if row is None:
            raise CustomerOrderWorkbookError(f"{sheet_name} 第 {row_number} 行不存在")
        self._set_cell(row, column, row_number, **payload)
        self.parts[path] = etree.tostring(
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

    def _highlight_style_index(
        self,
        base_style_index: int,
        *,
        fill_rgb: str,
    ) -> int:
        normalized_rgb = fill_rgb.upper()
        cache = getattr(self, "_highlight_style_cache", None)
        if cache is None:
            cache = {}
            self._highlight_style_cache = cache
        cache_key = (base_style_index, normalized_rgb)
        if cache_key in cache:
            return cache[cache_key]

        styles_path = "xl/styles.xml"
        root = etree.fromstring(self.parts[styles_path])
        fills = root.find("m:fills", NS)
        cell_xfs = root.find("m:cellXfs", NS)
        if fills is None or cell_xfs is None:
            raise CustomerOrderWorkbookError("工作簿样式表缺少 fills 或 cellXfs")

        fill_id: int | None = None
        for index, fill in enumerate(fills.findall("m:fill", NS)):
            pattern = fill.find("m:patternFill", NS)
            foreground = pattern.find("m:fgColor", NS) if pattern is not None else None
            if (
                pattern is not None
                and pattern.get("patternType") == "solid"
                and foreground is not None
                and foreground.get("rgb", "").upper() == normalized_rgb
            ):
                fill_id = index
                break
        if fill_id is None:
            fill_id = len(fills.findall("m:fill", NS))
            fill = etree.SubElement(fills, f"{{{MAIN_NS}}}fill")
            pattern = etree.SubElement(
                fill,
                f"{{{MAIN_NS}}}patternFill",
                patternType="solid",
            )
            etree.SubElement(
                pattern,
                f"{{{MAIN_NS}}}fgColor",
                rgb=normalized_rgb,
            )
            etree.SubElement(
                pattern,
                f"{{{MAIN_NS}}}bgColor",
                indexed="64",
            )
            fills.set("count", str(fill_id + 1))

        xfs = cell_xfs.findall("m:xf", NS)
        if not xfs:
            raise CustomerOrderWorkbookError("工作簿样式表没有可复用的单元格格式")
        if base_style_index < 0 or base_style_index >= len(xfs):
            base_style_index = 0
        source = xfs[base_style_index]
        if int(source.get("fillId", "0")) == fill_id:
            cache[cache_key] = base_style_index
            return base_style_index

        highlighted = deepcopy(source)
        highlighted.set("fillId", str(fill_id))
        highlighted.set("applyFill", "1")
        cell_xfs.append(highlighted)
        highlighted_index = len(xfs)
        cell_xfs.set("count", str(highlighted_index + 1))
        self.parts[styles_path] = etree.tostring(
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )
        cache[cache_key] = highlighted_index
        return highlighted_index

    def highlight_row(
        self,
        sheet_name: str,
        row_number: int,
        *,
        fill_rgb: str = "FFFFFF00",
    ) -> None:
        """Apply a visible fill to every materialized cell in one inserted row."""
        path = self.sheet_paths[sheet_name]
        root = etree.fromstring(self.parts[path])
        row = next(
            (
                item
                for item in root.findall("m:sheetData/m:row", NS)
                if int(item.get("r", "0")) == row_number
            ),
            None,
        )
        if row is None:
            raise CustomerOrderWorkbookError(f"{sheet_name} 第 {row_number} 行不存在")
        for cell in row.findall("m:c", NS):
            base_style_index = int(cell.get("s", "0"))
            cell.set(
                "s",
                str(
                    self._highlight_style_index(
                        base_style_index,
                        fill_rgb=fill_rgb,
                    )
                ),
            )
        self.parts[path] = etree.tostring(
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

    def insert_row_at(
        self,
        sheet_name: str,
        *,
        insert_row: int,
        values: dict[str, dict[str, Any]],
        reference_row_number: int | None = None,
    ) -> int:
        path = self.sheet_paths[sheet_name]
        root = etree.fromstring(self.parts[path])
        sheet_data = root.find("m:sheetData", NS)
        rows = sheet_data.findall("m:row", NS)
        if reference_row_number is None:
            reference_row = next(
                (row for row in reversed(rows) if int(row.get("r", "0")) < insert_row),
                None,
            )
        else:
            reference_row = next(
                (row for row in rows if int(row.get("r", "0")) == reference_row_number),
                None,
            )
        if reference_row is None:
            raise CustomerOrderWorkbookError(f"{sheet_name} 无可复用的数据行样式")
        reference_row = deepcopy(reference_row)

        for row in reversed(rows):
            row_number = int(row.get("r", "0"))
            if row_number < insert_row:
                continue
            row.set("r", str(row_number + 1))
            for cell in row.findall("m:c", NS):
                cell.set("r", self._shift_cell_reference(cell.get("r", ""), insert_row))

        for formula_node in root.findall(".//m:f", NS):
            if formula_node.text:
                formula_node.text = self._shift_formula(formula_node.text, insert_row)

        for node in root.iter():
            for attribute in ("ref", "sqref"):
                ref_value = node.get(attribute)
                if ref_value and re.search(r"[A-Z]+\d", ref_value):
                    node.set(attribute, self._shift_range_reference(ref_value, insert_row))

        new_row = reference_row
        new_row.set("r", str(insert_row))
        new_row.set("ht", "12")
        new_row.set("customHeight", "1")
        for cell in new_row.findall("m:c", NS):
            match = CELL_REF_PATTERN.match(cell.get("r", ""))
            if match:
                cell.set("r", f"{match.group(1)}{insert_row}")
            self._clear_cell(cell)
        for column, payload in values.items():
            self._set_cell(new_row, column, insert_row, **payload)

        insertion_index = next(
            (index for index, row in enumerate(sheet_data.findall("m:row", NS)) if int(row.get("r", "0")) > insert_row),
            len(sheet_data),
        )
        sheet_data.insert(insertion_index, new_row)

        self.parts[path] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        return insert_row

    def insert_row(
        self,
        sheet_name: str,
        *,
        marker_column: str,
        marker: str,
        values: dict[str, dict[str, Any]],
    ) -> int:
        return self.insert_row_at(
            sheet_name,
            insert_row=self._find_marker_row(sheet_name, marker_column, marker),
            values=values,
        )

    def set_recalculation(self) -> None:
        # Row insertion invalidates the workbook's cached calculation chain.
        # Excel may refuse to open an otherwise valid OOXML package when the
        # chain still points at cells that moved or were replaced, so remove
        # the optional part and let Excel rebuild it on first open.
        self.parts.pop("xl/calcChain.xml", None)
        self.part_info.pop("xl/calcChain.xml", None)

        relationships_path = "xl/_rels/workbook.xml.rels"
        relationships = etree.fromstring(self.parts[relationships_path])
        for relationship in list(relationships):
            if (
                relationship.get("Type", "").endswith("/calcChain")
                or relationship.get("Target", "").replace("\\", "/").endswith("calcChain.xml")
            ):
                relationships.remove(relationship)
        self.parts[relationships_path] = etree.tostring(
            relationships,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        content_types_path = "[Content_Types].xml"
        content_types = etree.fromstring(self.parts[content_types_path])
        for override in list(content_types):
            if override.get("PartName", "").endswith("/xl/calcChain.xml"):
                content_types.remove(override)
        self.parts[content_types_path] = etree.tostring(
            content_types,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        root = etree.fromstring(self.parts["xl/workbook.xml"])
        calc_pr = root.find("m:calcPr", NS)
        if calc_pr is None:
            calc_pr = etree.SubElement(root, f"{{{MAIN_NS}}}calcPr")
        calc_pr.set("fullCalcOnLoad", "1")
        calc_pr.set("forceFullCalc", "1")
        calc_pr.set("calcMode", "auto")
        self.parts["xl/workbook.xml"] = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )

    def to_bytes(self) -> bytes:
        output = BytesIO()
        with ZipFile(output, "w", compression=ZIP_DEFLATED, compresslevel=6) as archive:
            for name, payload in self.parts.items():
                archive.writestr(self.part_info[name], payload)
        return output.getvalue()


def _make_issue(severity: str, code: str, field: str, message: str) -> dict[str, Any]:
    return {"severity": severity, "code": code, "field": field, "message": message}


def _requires_po_number(parsed: ParsedPoLine) -> bool:
    customer_name = _clean_text(parsed.values.get("customer_name")).upper()
    return (
        parsed.input_template.startswith("BUZZBEE_WALMART_")
        or customer_name in {"WMC", "WMU", "WALMART"}
        or customer_name.startswith("WALMART ")
    )


def _build_preview_rows(
    parsed_lines: list[ParsedPoLine],
    schedule_index: dict[str, ScheduleLookup],
    received_date: str,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for index, parsed in enumerate(parsed_lines, start=1):
        values = parsed.values
        product_no = _clean_identifier(values.get("product_no"))
        schedule = schedule_index.get(product_no, ScheduleLookup(existing_rows=[]))
        quantity = _decimal(values.get("quantity"))
        units_per_carton = _decimal(values.get("units_per_carton"))
        carton_count = (
            quantity / units_per_carton
            if quantity is not None and units_per_carton is not None and units_per_carton > 0
            else None
        )
        unit_price = schedule.unit_price_hkd
        amount = quantity * unit_price if quantity is not None and unit_price is not None else None
        inspection_date = _clean_text(values.get("inspection_date"))
        inspection_raw = _clean_text(values.get("inspection_raw"))
        requested_ship_date = _clean_text(values.get("requested_ship_date"))
        line_q = inspection_date or requested_ship_date
        customer_q = inspection_date or inspection_raw
        issues: list[dict[str, Any]] = []

        required_fields = {
            "contract_no": "Contract No.",
            "customer_name": "客户名称",
            "product_no": "产品编号",
            "product_name_en": "产品名称",
            "quantity": "数量",
            "units_per_carton": "装箱数",
            "standard": "国家标准",
            "requested_ship_date": "客要求走货期",
        }
        if _requires_po_number(parsed) and not _clean_text(values.get("po_no")):
            issues.append(
                _make_issue(
                    "blocked",
                    "missing_required_field",
                    "po_no",
                    "WM 客 P/O# 未能从 PO 提取",
                )
            )
        for field, label in required_fields.items():
            if not _clean_text(values.get(field)):
                issues.append(_make_issue("blocked", "missing_required_field", field, f"{label} 未能从 PO 提取"))
        if not schedule.product_name_zh:
            issues.append(
                _make_issue(
                    "blocked",
                    "missing_product_name_zh",
                    "product_name_zh",
                    f"产品 {product_no or '—'} 在当前排期中没有可复用的中文名称",
                )
            )
        if unit_price is None:
            issues.append(
                _make_issue(
                    "blocked",
                    "missing_unit_price",
                    "unit_price_hkd",
                    f"产品 {product_no or '—'} 在当前排期中没有可确认的 HKD 单价",
                )
            )
        if carton_count is not None and carton_count != carton_count.to_integral_value():
            issues.append(
                _make_issue(
                    "blocked",
                    "partial_carton",
                    "carton_count",
                    "数量不能被装箱数整除，系统不会静默向上取整",
                )
            )
        if not values.get("packaging"):
            issues.append(
                _make_issue("warning", "missing_packaging", "packaging", "未从 PO 识别到包装编号")
            )

        po_no = _clean_text(values.get("po_no"))
        contract_no = _clean_identifier(values.get("contract_no"))
        for existing in schedule.existing_rows or []:
            if (
                existing["po_no"] == po_no
                and existing["contract_no"] == contract_no
                and existing["product_no"] == product_no
            ):
                issues.append(
                    _make_issue(
                        "blocked",
                        "existing_order_line",
                        "po_no",
                        (
                            f"当前排期接单表第 {existing['row']} 行已有相同 PO/合同/产品；"
                            "测试阶段可人工确认后重复导入"
                        ),
                    )
                )
                existing_date = existing.get("requested_ship_date", "")
                if existing_date and requested_ship_date and existing_date != requested_ship_date:
                    issues.append(
                        _make_issue(
                            "warning",
                            "requested_ship_date_conflict",
                            "requested_ship_date",
                            f"PO 日期为 {requested_ship_date}，现有排期日期为 {existing_date}；输出保留 PO 原值",
                        )
                    )
                break

        row_id = f"buzzbee-{contract_no or index}-{product_no or index}-{index}"
        for issue in issues:
            can_skip = (
                issue["severity"] == "blocked"
                and issue["code"] in SKIPPABLE_BLOCKER_LABELS
            )
            issue["can_skip"] = can_skip
            issue["skip_key"] = (
                f"{row_id}|{issue['code']}|{issue['field']}" if can_skip else ""
            )
            issue["skip_label"] = (
                SKIPPABLE_BLOCKER_LABELS[issue["code"]] if can_skip else ""
            )

        status = (
            "blocked"
            if any(issue["severity"] == "blocked" for issue in issues)
            else "warning"
            if issues
            else "valid"
        )
        lineage = {
            "received_date": "业务登记 · 来单日期",
            "po_no": parsed.lineage.get("po_no", ""),
            "contract_no": parsed.lineage.get("contract_no", ""),
            "customer_country": f"{parsed.lineage.get('customer_name', '')} / {parsed.lineage.get('country', '')}",
            "product_no": parsed.lineage.get("product_no", ""),
            "product_name_zh": f"当前客户排期 · 产品 {product_no} 最近有效中文名称",
            "product_name_en": parsed.lineage.get("product_name_en", ""),
            "quantity": parsed.lineage.get("quantity", ""),
            "units_per_carton": parsed.lineage.get("units_per_carton", ""),
            "carton_count": "系统计算 · 数量 ÷ 装箱数",
            "standard": parsed.lineage.get("standard", ""),
            "unit_price_hkd": f"当前客户排期 · 产品 {product_no} 最近有效 HKD 单价",
            "amount_hkd": "系统计算 · 数量 × 单价HK",
            "packaging": parsed.lineage.get("packaging", ""),
            "line_q": parsed.lineage.get("inspection_date") or parsed.lineage.get("requested_ship_date", ""),
            "customer_q": parsed.lineage.get("inspection_date") or parsed.lineage.get("requested_ship_date", ""),
            "requested_ship_date": parsed.lineage.get("requested_ship_date", ""),
        }
        rows.append(
            {
                "id": row_id,
                "status": status,
                "status_label": {"valid": "有效", "warning": "警告", "blocked": "阻断"}[status],
                "received_date": received_date,
                "po_no": po_no,
                "contract_no": contract_no,
                "customer_country": f"{values.get('customer_name', '')} / {values.get('country', '')}".strip(" /"),
                "customer_name": _clean_text(values.get("customer_name")),
                "country": _clean_text(values.get("country")),
                "product_no": product_no,
                "product_name_zh": schedule.product_name_zh,
                "product_name_en": _clean_text(values.get("product_name_en")),
                "quantity": _decimal_text(quantity),
                "units_per_carton": _decimal_text(units_per_carton),
                "carton_count": _decimal_text(carton_count),
                "standard": _clean_text(values.get("standard")),
                "unit_price_hkd": _decimal_text(unit_price, 6),
                "amount_hkd": _decimal_text(amount, 4),
                "packaging": _clean_text(values.get("packaging")),
                "line_q": line_q,
                "customer_q": customer_q,
                "requested_ship_date": requested_ship_date,
                "input_template": parsed.input_template,
                "target_template": TARGET_TEMPLATE,
                "source_po_file_name": parsed.source_file_name,
                "lineage": lineage,
                "issues": issues,
            }
        )
    return rows


def _output_file_name(schedule_file_name: str) -> str:
    file_name = re.split(r"[\\/]", schedule_file_name)[-1].strip()
    return file_name or "BuzzBee生产排期表.xlsx"


def _mark_batch_duplicates(rows: list[dict[str, Any]]) -> None:
    seen: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    for row in rows:
        identity = (
            row["po_no"],
            row["contract_no"],
            row["product_no"],
            row["requested_ship_date"],
        )
        previous = seen.get(identity)
        if previous is None:
            seen[identity] = row
            continue
        issue = _make_issue(
            "blocked",
            "duplicate_batch_order_line",
            "po_no",
            (
                f"本批次与 {previous['source_po_file_name']} 存在相同订单行；"
                "测试阶段可人工确认后分别写入"
            ),
        )
        issue.update({
            "can_skip": True,
            "skip_key": f"{row['id']}|duplicate_batch_order_line|po_no",
            "skip_label": "测试阶段确认重复导入本批相同订单行",
        })
        row["issues"].append(issue)
        row["status"] = "blocked"
        row["status_label"] = "阻断"


def create_buzzbee_batch_preview(
    *,
    factory_id: str,
    received_date: str,
    po_files: list[tuple[str, bytes]],
    schedule_file_name: str,
    schedule_content: bytes,
) -> dict[str, Any]:
    if not po_files:
        raise CustomerOrderWorkbookError("请至少上传一份 PO 文件")
    if len(po_files) > MAX_BATCH_PO_FILES:
        raise CustomerOrderWorkbookError(
            f"单批最多上传 {MAX_BATCH_PO_FILES} 份 PO 文件"
        )
    if sum(len(content) for _, content in po_files) > MAX_BATCH_PO_BYTES:
        raise CustomerOrderWorkbookError("本批 PO 文件合计超过 80MB 限制")
    oversized = [name for name, content in po_files if len(content) > MAX_PO_BYTES]
    if oversized:
        raise CustomerOrderWorkbookError(
            f"PO 文件超过 12MB 限制：{', '.join(oversized)}"
        )
    if len(schedule_content) > MAX_SCHEDULE_BYTES:
        raise CustomerOrderWorkbookError("客户排期超过 35MB 限制")
    if "(印尼)" in schedule_file_name:
        raise CustomerOrderWorkbookError("印尼排期不属于当前 BuzzBee 映射范围")
    try:
        normalized_received_date = date.fromisoformat(received_date).isoformat()
    except ValueError as exc:
        raise CustomerOrderWorkbookError("来单日期必须是 YYYY-MM-DD") from exc

    parsed_lines: list[ParsedPoLine] = []
    for po_file_name, po_content in po_files:
        try:
            parsed_lines.extend(parse_po(po_file_name, po_content))
        except CustomerOrderWorkbookError as exc:
            raise CustomerOrderWorkbookError(f"{po_file_name}：{exc}") from exc
    plain_schedule, encrypted = _decrypt_schedule(schedule_content)
    schedule = OoxmlSchedule(plain_schedule)
    rows = _build_preview_rows(parsed_lines, schedule.build_schedule_index(), normalized_received_date)
    for row in rows:
        row["item_sheet_name"] = _item_sheet_name(row)
    _mark_batch_duplicates(rows)
    summary = {
        "total": len(rows),
        "valid": sum(row["status"] == "valid" for row in rows),
        "warning": sum(row["status"] == "warning" for row in rows),
        "blocked": sum(row["status"] == "blocked" for row in rows),
    }
    input_templates = sorted({row["input_template"] for row in rows})
    warnings = [
        "输出将更新“接单表”“正单评审表”及每行对应的子弹枪/水枪 ITEM 表；ITEM 新增订单行，存在同货号备料单时，其 H 列按本合同数量扣减；不自动补 3000，也不新增样板或备料行。",
        "下载文件保持原排期的 2026 打开密码。" if encrypted else "上传文件未加密；下载文件将使用 BuzzBee 2026 打开密码。",
    ]
    if len(po_files) > 1:
        warnings.append(
            f"本批共 {len(po_files)} 份 PO 文件，确认后合并写入同一份客户排期。"
        )
    po_file_names = [name for name, _ in po_files]
    po_hashes = [sha256(content).hexdigest() for _, content in po_files]
    combined_po_hash = (
        po_hashes[0]
        if len(po_hashes) == 1
        else sha256(
            "\n".join(
                f"{name}:{digest}"
                for name, digest in zip(po_file_names, po_hashes, strict=True)
            ).encode("utf-8")
        ).hexdigest()
    )
    return {
        "preview_schema_version": PREVIEW_SCHEMA_VERSION,
        "customer_code": "buzzbee",
        "factory_id": factory_id,
        "po_file_name": po_file_names[0] if len(po_file_names) == 1 else f"{len(po_file_names)}个PO文件",
        "po_file_names": po_file_names,
        "po_file_count": len(po_file_names),
        "schedule_file_name": schedule_file_name,
        "source_po_sha256": combined_po_hash,
        "source_po_sha256s": po_hashes,
        "source_schedule_sha256": sha256(schedule_content).hexdigest(),
        "input_template": ", ".join(input_templates),
        "target_template": TARGET_TEMPLATE,
        "output_file_name": _output_file_name(schedule_file_name),
        "summary": summary,
        "rows": rows,
        "warnings": warnings,
    }


def create_buzzbee_preview(
    *,
    factory_id: str,
    received_date: str,
    po_file_name: str,
    po_content: bytes,
    schedule_file_name: str,
    schedule_content: bytes,
) -> dict[str, Any]:
    return create_buzzbee_batch_preview(
        factory_id=factory_id,
        received_date=received_date,
        po_files=[(po_file_name, po_content)],
        schedule_file_name=schedule_file_name,
        schedule_content=schedule_content,
    )


def _order_row_values(row: dict[str, Any]) -> dict[str, dict[str, Any]]:
    quantity = Decimal(row["quantity"])
    packing = Decimal(row["units_per_carton"])
    cartons = Decimal(row["carton_count"])
    price = _decimal(row["unit_price_hkd"])
    amount = _decimal(row["amount_hkd"])
    us_standard = row["standard"] == "美国标准"
    discount_price = price * Decimal("0.98") if price is not None else None
    discount_amount = amount * Decimal("0.98") if amount is not None else None
    return {
        "B": {"value": _excel_serial(row["received_date"])},
        "C": {"value": row["po_no"], "inline": True},
        "D": {"value": row["contract_no"], "inline": True},
        "E": {"value": row["customer_name"], "inline": True},
        "F": {"value": row["product_no"], "inline": True},
        "G": {"value": row["product_name_zh"], "inline": True},
        "H": {"value": row["product_name_en"], "inline": True},
        "I": {"value": quantity},
        "J": {"value": packing},
        "K": {"formula": f"I{{row}}/J{{row}}", "cached": cartons},
        "L": {"value": row["standard"], "inline": True},
        "M": {"value": price} if price is not None else {},
        "N": (
            {"formula": "M{row}*0.98", "cached": discount_price}
            if us_standard and price is not None
            else {}
        ),
        "O": (
            {"formula": "I{row}*M{row}", "cached": amount}
            if price is not None
            else {}
        ),
        "P": (
            {"formula": "O{row}*0.98", "cached": discount_amount}
            if us_standard and amount is not None
            else {}
        ),
        "Q": {"value": row["packaging"], "inline": True},
        "S": {"value": _excel_serial(row["requested_ship_date"])},
    }


def _is_water_product(row: dict[str, Any]) -> bool:
    combined = f"{row['product_name_zh']} {row['product_name_en']}".upper()
    bullet_keywords = ("子弹", "镖弹", "飞镖", "软弹", "泡沫弹", "DART", "BULLET", "FOAM", "REFILL")
    if any(keyword in combined for keyword in bullet_keywords):
        return False
    return any(keyword in combined for keyword in ("水枪", "水炮", "喷水", "SQUIRT", "WATER", "SOAKER", "SPLASH"))


def _item_sheet_name(row: dict[str, Any]) -> str:
    return ITEM_SHEET_WATER if _is_water_product(row) else ITEM_SHEET_BULLET


def _item_boundary_row(rows: dict[int, dict[str, str]]) -> int:
    for row_number in sorted(rows):
        if any(
            keyword in value
            for value in rows[row_number].values()
            for keyword in ITEM_BOUNDARY_KEYWORDS
        ):
            return row_number
    return max(rows, default=4) + 1


def _item_matches(stored_item_no: Any, base_item_no: str) -> bool:
    stored = _clean_identifier(stored_item_no)
    return stored == base_item_no or stored.startswith(f"{base_item_no}-")


def _item_contract_template_row(
    rows: dict[int, dict[str, str]],
    boundary_row: int,
) -> int:
    candidates: list[tuple[Decimal, int]] = []
    for row_number, values in rows.items():
        if row_number <= 3 or row_number >= boundary_row:
            continue
        if not (
            _clean_identifier(values.get("C"))
            and _clean_text(values.get("D"))
            and _clean_identifier(values.get("E"))
            and _decimal(values.get("H")) is not None
        ):
            continue
        candidates.append((_decimal(values.get("A")) or Decimal("0"), row_number))
    if candidates:
        return max(candidates)[1]
    fallback = next(
        (
            row_number
            for row_number, values in sorted(rows.items())
            if 3 < row_number < boundary_row and _clean_identifier(values.get("E"))
        ),
        None,
    )
    if fallback is None:
        raise CustomerOrderWorkbookError("ITEM 表活动区没有可复用的订单行样式")
    return fallback


def _item_oqf_no(received_date: str, item_no: str) -> str:
    parsed = date.fromisoformat(received_date)
    return f"{parsed.year}-{parsed.month}-{parsed.day}-{item_no}"


def _item_date_payload(value: str) -> dict[str, Any]:
    if not value:
        return {}
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        return {"value": _excel_serial(value)}
    return {"value": value, "inline": True}


def _item_row_values(
    row: dict[str, Any],
    *,
    item_no: str,
    oqf_no: str,
) -> dict[str, dict[str, Any]]:
    packing = _decimal(row["units_per_carton"])
    return {
        "A": {"value": _excel_serial(row["received_date"])},
        "B": {"value": oqf_no, "inline": True},
        "C": {"value": row["contract_no"], "inline": True},
        "D": {"value": row["customer_name"], "inline": True},
        "E": {"value": item_no, "inline": True},
        "F": {"value": row["product_name_zh"], "inline": True},
        "G": {"value": row["product_name_en"], "inline": True},
        "H": {"value": Decimal(row["quantity"])},
        "I": {"value": packing} if packing is not None else {},
        "K": {"value": row["packaging"], "inline": True},
        "L": {"value": "系统", "inline": True},
        "M": _item_date_payload(row["line_q"]),
        "N": _item_date_payload(row["customer_q"]),
        "O": _item_date_payload(row["requested_ship_date"]),
        "Q": {"value": row["po_no"], "inline": True},
    }


def _write_item_order_row(
    workbook: OoxmlSchedule,
    row: dict[str, Any],
) -> dict[str, Any]:
    sheet_name = _item_sheet_name(row)
    rows = workbook.read_rows(sheet_name)
    boundary_row = _item_boundary_row(rows)
    base_item_no = _clean_identifier(row["product_no"])
    item_rows = [
        row_number
        for row_number, values in sorted(rows.items())
        if 3 < row_number < boundary_row and _item_matches(values.get("E"), base_item_no)
    ]
    template_row = _item_contract_template_row(rows, boundary_row)
    preparation_row: int | None = None

    if item_rows:
        preparation_row = next(
            (
                row_number
                for row_number in item_rows
                if "备料单" in f"{rows[row_number].get('L', '')}{rows[row_number].get('M', '')}"
            ),
            None,
        )
        anchor_row = preparation_row or item_rows[0]
        resolved_item_no = _clean_identifier(rows[anchor_row].get("E")) or base_item_no
        oqf_no = _clean_text(rows[anchor_row].get("B")) or _item_oqf_no(
            row["received_date"],
            resolved_item_no,
        )
        insert_row = item_rows[0]
    else:
        resolved_item_no = base_item_no
        oqf_no = _item_oqf_no(row["received_date"], resolved_item_no)
        divider_row = max(
            (
                row_number
                for row_number, values in rows.items()
                if row_number < boundary_row and _clean_identifier(values.get("E"))
            ),
            default=4,
        ) + 1
        insert_row = divider_row
        if divider_row > 5:
            workbook.insert_row_at(
                sheet_name,
                insert_row=divider_row,
                values={},
                reference_row_number=template_row,
            )
            insert_row += 1

    workbook.insert_row_at(
        sheet_name,
        insert_row=insert_row,
        values=_item_row_values(row, item_no=resolved_item_no, oqf_no=oqf_no),
        reference_row_number=template_row,
    )
    if preparation_row is not None:
        shifted_preparation_row = preparation_row + 1
        shifted_rows = workbook.read_rows(sheet_name)
        preparation_formula = workbook.read_cell_formula(
            sheet_name,
            shifted_preparation_row,
            "H",
        ).strip()
        preparation_quantity = _decimal(
            shifted_rows.get(shifted_preparation_row, {}).get("H")
        )
        contract_quantity = Decimal(row["quantity"])
        formula_base = preparation_formula or _decimal_text(preparation_quantity)
        if not formula_base:
            raise CustomerOrderWorkbookError(
                f"{sheet_name} 第 {shifted_preparation_row} 行备料数量为空，无法扣减合同数量"
            )
        workbook.set_cell(
            sheet_name,
            shifted_preparation_row,
            "H",
            formula=f"{formula_base}-{_decimal_text(contract_quantity)}",
            cached=(
                preparation_quantity - contract_quantity
                if preparation_quantity is not None
                else None
            ),
        )
    return {
        "sheet_name": sheet_name,
        "inserted_row": insert_row,
        "item_no": resolved_item_no,
        "oqf_no": oqf_no,
    }


def _review_row_values(row: dict[str, Any]) -> dict[str, dict[str, Any]]:
    quantity = Decimal(row["quantity"])
    packing = Decimal(row["units_per_carton"])
    cartons = Decimal(row["carton_count"])
    price = _decimal(row["unit_price_hkd"])
    amount = _decimal(row["amount_hkd"])
    us_standard = row["standard"] == "美国标准"
    discount_price = price * Decimal("0.98") if price is not None else None
    discount_amount = amount * Decimal("0.98") if amount is not None else None
    values: dict[str, dict[str, Any]] = {
        "B": {"value": _excel_serial(row["received_date"])},
        "C": {"value": row["po_no"], "inline": True},
        "D": {"value": row["contract_no"], "inline": True},
        "E": {"value": row["customer_name"], "inline": True},
        "F": {"value": row["product_no"], "inline": True},
        "G": {"value": row["product_name_zh"], "inline": True},
        "H": {"value": row["product_name_en"], "inline": True},
        "I": {"value": quantity},
        "J": {"value": packing},
        "K": {"formula": "I{row}/J{row}", "cached": cartons},
        "L": {"value": _excel_serial(row["line_q"])} if row["line_q"] else {},
        "M": (
            {"value": _excel_serial(row["customer_q"])}
            if re.fullmatch(r"\d{4}-\d{2}-\d{2}", row["customer_q"])
            else {"value": row["customer_q"], "inline": True}
            if row["customer_q"]
            else {}
        ),
        "N": {"value": row["standard"], "inline": True},
        "Q": {"value": "已上", "inline": True},
        "R": {"value": price} if price is not None else {},
        "S": (
            {"formula": "R{row}*0.98", "cached": discount_price}
            if us_standard and price is not None
            else {}
        ),
        "T": (
            {"formula": "I{row}*R{row}", "cached": amount}
            if price is not None
            else {}
        ),
        "U": (
            {"formula": "T{row}*0.98", "cached": discount_amount}
            if us_standard and amount is not None
            else {}
        ),
        "V": {"value": row["packaging"], "inline": True},
        "W": {"value": _excel_serial(row["requested_ship_date"])},
    }
    return values


def _materialize_formulas(values: dict[str, dict[str, Any]], row_number: int) -> None:
    for payload in values.values():
        if payload.get("formula"):
            payload["formula"] = payload["formula"].format(row=row_number)


def export_buzzbee_batch_schedule(
    *,
    factory_id: str,
    received_date: str,
    po_files: list[tuple[str, bytes]],
    schedule_file_name: str,
    schedule_content: bytes,
    skipped_issue_keys: set[str] | None = None,
    manual_overrides: list[dict[str, str]] | None = None,
) -> tuple[bytes, str, dict[str, Any]]:
    preview = create_buzzbee_batch_preview(
        factory_id=factory_id,
        received_date=received_date,
        po_files=po_files,
        schedule_file_name=schedule_file_name,
        schedule_content=schedule_content,
    )
    decorate_manual_resolution_policy(preview)
    apply_overrides_to_preview(preview, manual_overrides or [])
    requested_skips = skipped_issue_keys or set()
    available_skips = {
        issue["skip_key"]
        for row in preview["rows"]
        for issue in row["issues"]
        if issue.get("skip_key")
    }
    invalid_skips = requested_skips - available_skips
    if invalid_skips:
        raise CustomerOrderWorkbookError(
            "所选跳过项已失效或不允许跳过，请重新解析后再确认"
        )
    remaining_blockers = [
        issue
        for row in preview["rows"]
        for issue in row["issues"]
        if issue["severity"] == "blocked"
        and issue["skip_key"] not in requested_skips
    ]
    if remaining_blockers:
        blocked_messages = [
            issue["message"] for issue in remaining_blockers
        ]
        raise CustomerOrderWorkbookError("仍有阻断项：" + "；".join(blocked_messages))

    plain_schedule, _ = _decrypt_schedule(schedule_content)
    workbook = OoxmlSchedule(plain_schedule)
    for row in preview["rows"]:
        order_values = _order_row_values(row)
        order_insert_row = workbook._find_marker_row("接单表", "M", "合计")
        _materialize_formulas(order_values, order_insert_row)
        workbook.insert_row(
            "接单表",
            marker_column="M",
            marker="合计",
            values=order_values,
        )

        marker = "水枪合计" if _is_water_product(row) else "子弹枪合计"
        review_values = _review_row_values(row)
        review_insert_row = workbook._find_marker_row("正单评审表", "R", marker)
        _materialize_formulas(review_values, review_insert_row)
        workbook.insert_row(
            "正单评审表",
            marker_column="R",
            marker=marker,
            values=review_values,
        )

    for row in reversed(preview["rows"]):
        _write_item_order_row(workbook, row)

    workbook.set_recalculation()
    output = _encrypt_schedule(workbook.to_bytes())
    return output, preview["output_file_name"], preview


def export_buzzbee_schedule(
    *,
    factory_id: str,
    received_date: str,
    po_file_name: str,
    po_content: bytes,
    schedule_file_name: str,
    schedule_content: bytes,
    skipped_issue_keys: set[str] | None = None,
    manual_overrides: list[dict[str, str]] | None = None,
) -> tuple[bytes, str, dict[str, Any]]:
    return export_buzzbee_batch_schedule(
        factory_id=factory_id,
        received_date=received_date,
        po_files=[(po_file_name, po_content)],
        schedule_file_name=schedule_file_name,
        schedule_content=schedule_content,
        skipped_issue_keys=skipped_issue_keys,
        manual_overrides=manual_overrides,
    )
