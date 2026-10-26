"""Recover complete native front-label blocks with unambiguous current-row identity.

This is a structured print layout, not the generic source-cell fallback. It only
accepts a complete one-to-one match using at least two explicit identity columns.
"""
import json
import re
from collections import defaultdict
from openpyxl.utils.cell import coordinate_to_tuple


def _key(text):
    return re.sub(r'[^a-z0-9]','',text.lower()).replace('facillity','facility')


def linked(source):
    """Recover repeated print faces whose cached cells explicitly cite each row."""
    from app.services import carton_mark_flexible as flow
    rows=source.get('detail_rows',[]);links=source.get('native_links',{})
    if not rows or not links:return None
    cells={c.id:c for c in source['cells']};row_for={key:r['id'] for r in rows for key in r['cell_ids']}
    notes={n.location for n in source['notes']};limits=flow.constraints(source)
    groups=defaultdict(list)
    for block in source.get('blocks',[]):groups[(block['sheet'],block['start_row'])].append(block)
    records={}
    def select(cell):
        key=links.get(cell.id,cell.id)
        return dict(source_id=key,value=cells[key].text)
    for blocks in groups.values():
        panels=[];anchors=set();identities=set()
        for block in blocks:
            grouped=defaultdict(list)
            for key in block['cell_ids']:
                cell=cells[key];_,co=cell.location.rsplit('!',1);r,col=coordinate_to_tuple(co)
                if cell.location not in notes:grouped[r].append((col,cell))
            fields=[];bars=[]
            for _,entries in sorted(grouped.items()):
                entries.sort(key=lambda value:value[0]);label=entries[0][1]
                if re.search(r'logo|shipping\s*mark|remark|印刷|注意|example',label.text,re.I):continue
                if len(entries)==1:
                    if re.fullmatch(r'MADE\s+IN\s+.+',label.text,re.I):fields.append(dict(label=None,parts=[select(label)],role='normal'))
                    continue
                # A literal native label plus its displayed cells forms a print
                # field. Captions, template zeros and errors cannot supply rows.
                if not re.search(r'[a-zA-Z]',label.text) or re.search(r'#N/A|#REF!',label.text):continue
                parts=[cell for _,cell in entries[1:] if not re.search(r'#N/A|#REF!',cell.text)]
                if not parts:continue
                for cell in parts:
                    target=links.get(cell.id)
                    if target in row_for:anchors.add(row_for[target]);identities.add(target)
                if re.search(r'barcode|^\s*(?:UPC|GTIN|EAN)',label.text,re.I) and len(parts)==1 and re.fullmatch(r'\d{12,14}',parts[0].text):
                    bars.append(dict(source=select(parts[0]),kind={12:'UPCA',13:'EAN13',14:'ITF14'}[len(parts[0].text)],width_mm=70.,height_mm=35.,note_ids=[]))
                else:fields.append(dict(label=dict(source_id=label.id,value=label.text),parts=[select(c) for c in parts],role='normal'))
            if fields or bars or limits['logo']:
                panels.append(dict(face=block['face'],copies=block.get('copies') or source.get('face_copies',{}).get(block['sheet'],{}).get(block['face'],2),logo=limits['logo'],fields=fields,barcodes=bars,symbols=[]))
        if len(anchors)!=1 or len(identities)<2:continue
        key=next(iter(anchors))
        if key in records or len({p['face'] for p in panels})!=len(panels):return None
        records[key]=dict(row_id=key,panels=panels)
    if set(records)!={r['id'] for r in rows}:return None
    try:plan=flow.validate(json.dumps(dict(template_id='source-flow-v1',paper='A3',orientation='landscape',records=[records[r['id']] for r in rows]),ensure_ascii=False),source,review_mode=True)
    except flow.pilot.PilotError:return None
    plan._warnings.append('完整箱面按原稿标签、单元格显示值及其明确的当前行引用自动恢复；未将辅助表内部物料号、总数量等替换为印刷字段。')
    return plan


def table_template(source):
    """Bind one explicit two-face example to a complete native detail table."""
    from app.services import carton_mark_flexible as flow
    rows=source.get('detail_rows',[]);blocks=source.get('blocks',[])
    groups=defaultdict(list)
    for block in blocks:groups[(block['sheet'],block['start_row'])].append(block)
    if not rows or len(groups)!=1:return None
    pair=next(iter(groups.values()))
    if len(pair)!=2 or {b['face'] for b in pair}!={'front','side'}:return None
    cells={c.id:c for c in source['cells']};locations={c.location:c for c in source['cells']}
    notes={n.location for n in source['notes']};first=rows[0]
    if len({(r['id'].rsplit('!row:',1)[0],r['header_row']) for r in rows})!=1:return None
    aliases={'orderno':'po','ordernumber':'po','ponumber':'po','pono':'po','item':'itemno','itemnumber':'itemno',
             'newdescription':'description','qtypcs':'qty','quantity':'qty','ctnno':'cartonno','cartonnumber':'cartonno'}
    canonical=lambda text:aliases.get(_key(text),_key(text))
    from openpyxl.utils import get_column_letter
    headers={};columns={}
    for key in first['cell_ids']:
        cell=cells[key];sheet,co=cell.location.rsplit('!',1);_,col=coordinate_to_tuple(co)
        header=locations.get(f"{sheet}!{get_column_letter(col)}{first['header_row']}")
        if header:headers.setdefault(canonical(header.text),[]).append((col,cell));columns[col]=header
    bindings={};panels=[]
    for block in pair:
        grouped=defaultdict(list)
        for key in block['cell_ids']:
            cell=cells[key];_,co=cell.location.rsplit('!',1);r,col=coordinate_to_tuple(co)
            if cell.location not in notes:grouped[r].append((col,cell))
        fields=[];bars=[]
        for _,entries in sorted(grouped.items()):
            entries.sort(key=lambda entry:entry[0]);label=entries[0][1]
            if len(entries)==1:
                if re.fullmatch(r'MADE\s+IN\s+.+',label.text,re.I):fields.append((None,[label]))
                elif re.fullmatch(r'\d{12,14}',label.text):
                    possible=[(col,value) for name,items in headers.items() if name in {'barcode','gtin','upc','ean'} for col,value in items if flow.pilot.normalized(value.text)==flow.pilot.normalized(label.text)]
                    if len(possible)!=1:return None
                    bindings[label.id]=possible[0][0];bars.append(label)
                continue
            if re.search(r'logo|remark|example|印刷|注意',label.text,re.I):continue
            possible=headers.get(canonical(label.text),[])
            if len(possible)!=1:return None
            col,value=possible[0];parts=[c for _,c in entries[1:]]
            exact=[c for c in parts if flow.pilot.normalized(c.text)==flow.pilot.normalized(value.text)]
            if len(exact)!=1:return None
            bindings[exact[0].id]=col;fields.append((label,parts))
        panels.append((block,fields,bars))
    if len(set(bindings.values()))<4:return None
    records=[];limits=flow.constraints(source)
    for row in rows:
        current={coordinate_to_tuple(cells[key].location.rsplit('!',1)[1])[1]:cells[key] for key in row['cell_ids']}
        if set(bindings.values())-set(current):return None
        def selected(cell):
            native=current[bindings[cell.id]] if cell.id in bindings else cell
            return dict(source_id=native.id,value=native.text)
        output=[]
        for block,fields,bars in panels:
            output.append(dict(face=block['face'],copies=block.get('copies') or 2,logo=limits['logo'],symbols=[],
                fields=[dict(label=dict(source_id=label.id,value=label.text) if label else None,parts=[selected(c) for c in parts],role='normal') for label,parts in fields],
                barcodes=[dict(source=selected(c),kind={12:'UPCA',13:'EAN13',14:'ITF14'}.get(len(selected(c)['value']),'Code128'),width_mm=70.,height_mm=35.,note_ids=[]) for c in bars]))
        records.append(dict(row_id=row['id'],panels=output))
    try:plan=flow.validate(json.dumps(dict(template_id='source-flow-v1',paper='A3',orientation='landscape',records=records),ensure_ascii=False),source,review_mode=True)
    except flow.pilot.PilotError:return None
    plan._warnings.append('已按唯一正侧唛模板的原文标签与首行明细逐项一致关系绑定当前表列，分别生成每行；模板单位和分面保留，未复制示例数字。')
    return plan


def build(source):
    from app.services import carton_mark_flexible as flow
    rows=source.get('detail_rows',[])
    if not rows:return None
    cells={c.id:c for c in source['cells']};notes={n.location for n in source['notes']}
    locations={c.location:c for c in source['cells']};records={}
    identity_keys={'customerpono','customerpo','ponumber','pono','upc','barcode','exchangedestinationfacility'}
    def column_for(row,label):
        sh=row['id'].rsplit('!row:',1)[0];wanted=_key(label)
        if wanted in {'ponumber','pono'}:wanted='customerpono'
        aliases={'customerpo':'customerpono','itemdescription':'itemdescription'}
        wanted=aliases.get(wanted,wanted)
        candidates=[]
        from openpyxl.utils import get_column_letter
        for key in row['cell_ids']:
            cell=cells[key];_,co=cell.location.rsplit('!',1);_,col=coordinate_to_tuple(co)
            header=locations.get(f"{sh}!{get_column_letter(col)}{row['header_row']}")
            if header and aliases.get(_key(header.text),_key(header.text))==wanted:candidates.append(cell)
        return candidates[0] if len(candidates)==1 else None
    for block in source.get('blocks',[]):
        if block['face']!='front' or source.get('face_copies',{}).get(block['sheet'],{}).get('front') not in {1,2,4}:continue
        # Restrict recovery to a clearly declared single-face label layout.
        if any(b['sheet']==block['sheet'] and b['start_row']==block['start_row'] and b['face']=='side' for b in source['blocks']):continue
        grouped=defaultdict(list)
        for key in block['cell_ids']:
            cell=cells[key];_,co=cell.location.rsplit('!',1);r,col=coordinate_to_tuple(co)
            if cell.location in notes or cell.kind=='reference_cell':continue
            grouped[r].append((col,cell))
        fields=[];bars=[];pending=None
        for r,entries in sorted(grouped.items()):
            entries.sort(key=lambda v:v[0]);label=entries[0][1]
            if re.search(r'印刷|贴|注意|remark|shipping\s*mark|logo|example',label.text,re.I):continue
            if len(entries)==1:
                if _key(label.text) in {'upc','barcode','ean','gtin'}:pending=label;continue
                if pending and re.fullmatch(r'\d{12,14}',label.text):
                    bars.append((pending,label));pending=None;continue
                if re.fullmatch(r'MADE\s+IN\s+.+',label.text,re.I):fields.append((None,[label]))
                continue
            values=[c for _,c in entries[1:]]
            fields.append((label,values))
        candidates=[]
        for row in rows:
            identities=[];valid=True
            for label,parts in fields+[(label,[value]) for label,value in bars]:
                if not label or _key(label.text) not in identity_keys:continue
                target=column_for(row,label.text)
                if target is None or len(parts)!=1 or flow.pilot.normalized(parts[0].text)!=flow.pilot.normalized(target.text):valid=False;break
                identities.append(target.id)
            if valid and len(set(identities))>=2:candidates.append(row)
        if len(candidates)!=1:continue
        row=candidates[0]
        if row['id'] in records:return None
        def choice(cell):return dict(source_id=cell.id,value=cell.text)
        printed=[]
        for label,parts in fields:
            target=column_for(row,label.text) if label else None
            # Equal native identity values can use the actual current row ID.
            # Other literal print text (e.g. a shortened product name) stays as
            # the supplied print block, never rewritten from a catalog.
            if target and len(parts)==1 and flow.pilot.normalized(parts[0].text)==flow.pilot.normalized(target.text):parts=[target]
            printed.append(dict(label=choice(label) if label else None,parts=[choice(c) for c in parts],role='normal'))
        barcodes=[]
        for label,cell in bars:
            target=column_for(row,label.text)
            if target and flow.pilot.normalized(target.text)==flow.pilot.normalized(cell.text):cell=target
            barcodes.append(dict(source=choice(cell),kind={12:'UPCA',13:'EAN13',14:'ITF14'}[len(cell.text)],width_mm=70.,height_mm=35.,note_ids=[]))
        if not barcodes or len(printed)<4:continue
        records[row['id']]=dict(row_id=row['id'],panels=[dict(face='front',copies=source['face_copies'][block['sheet']]['front'],logo=flow.constraints(source)['logo'],fields=printed,barcodes=barcodes,symbols=[])])
    if set(records)!={r['id'] for r in rows}:return None
    try:
        plan=flow.validate(json.dumps(dict(template_id='source-flow-v1',paper='A3',orientation='landscape',records=[records[r['id']] for r in rows]),ensure_ascii=False),source,review_mode=True)
    except flow.pilot.PilotError:return None
    plan._warnings.append('已按原稿完整箱面区块与当前明细的PO、目的地/条码逐组唯一对应，自动恢复全部箱面；没有沿用模型漏行结果。')
    return plan
