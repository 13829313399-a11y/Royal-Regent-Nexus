from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from statistics import median
from typing import Any

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.carton_procurement import CartonOrder, CartonOrderLine


MAX_IMPORT_ROWS = 5_000
MAX_PREVIEW_ROWS = 500
# This value is also part of the delivery-import deduplication identity. Bump it
# whenever an OCR/parser change must reprocess files imported by an older build.
DELIVERY_IMPORT_PARSER_VERSION = "delivery-note-local-v7-dongkang"
SCHEDULE_IMPORT_PARSER_VERSION = "schedule-item-v1"
PACKAGING_TYPES = (
    "普通箱",
    "压线卡",
    "半亦箱",
    "全亦箱",
    "滑板纸",
    "展示盒",
    "内箱",
    "外箱",
    "平卡",
    "压卡",
    "啤卡",
    "刀卡",
    "卡纸",
    "天盖",
    "地盖",
    "纸箱",
)


WEEKLY_ALIASES = {
    "customer_po": {"客户po", "客户采购单号", "客户订单号", "customerpo"},
    "contract_no": {"reference", "合同号", "合同", "sc", "参考号"},
    "po_numbers": {"pono", "po", "po号", "客户po", "订单号"},
    "customer_name": {"客名国家", "客名", "客户", "国家", "customer"},
    "item_no": {"产品编号", "货号", "客货号", "itemno", "item"},
    "product_name": {"产品名称", "品名", "description", "productname"},
    "quantity": {"数量", "订单数量", "qty", "quantity"},
    "carton_rule": {"装箱", "包装", "装箱方式", "pack", "packing"},
    "inspection_window": {"验货期请提前准备好货物", "验货期", "查货期", "inspectiondate"},
}

DELIVERY_ALIASES = {
    "customer_po": {"客户po", "客户采购单号", "客户订单号", "customerpo"},
    "delivery_date": {"日期", "送货日期", "送货时间", "入库日期", "date"},
    "delivery_note_no": {"入库单号", "送货单号", "送货单", "dnno", "dn"},
    "contract_no": {"po", "合同号", "合同", "pono", "订单编号", "客户单号"},
    "item_no": {"货号", "产品编号", "itemno", "item", "客户料号"},
    "quantity": {"外箱", "送货数量", "数量", "入库数量", "qty", "quantity"},
    "paper_quality": {"纸质", "材质", "paperquality", "description"},
    "length": {"长", "l", "length"},
    "width": {"宽", "w", "width"},
    "height": {"高", "h", "height"},
    "unit_price": {"单价", "unitprice", "price"},
    "location": {"仓位", "库位", "location"},
    "destination": {"客户"},
    "packaging_type": {"名称"},
    "specification": {"规格"},
}


@dataclass(frozen=True)
class DeliveryOcrWord:
    text: str
    left: int
    top: int
    right: int
    bottom: int
    confidence: float

    @property
    def center_x(self) -> float:
        return (self.left + self.right) / 2

    @property
    def center_y(self) -> float:
        return (self.top + self.bottom) / 2


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", _text(value).lower())


def _identity(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", _text(value).lower())


def _dongkang_identity(value: Any) -> str:
    # Customer item numbers can be Chinese product names in the supplier file.
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", _text(value).casefold())


def _dimension_identity(value: Any) -> tuple[str, ...]:
    dimensions = re.findall(r"\d+(?:[.,]\d+)?", _text(value))
    if len(dimensions) < 2:
        return ()
    result: list[str] = []
    for dimension in dimensions[:3]:
        try:
            result.append(str(Decimal(dimension.replace(",", ".")).normalize()))
        except InvalidOperation:
            return ()
    return tuple(result)


def _number(value: Any) -> Decimal | None:
    if value in (None, ""):
        return None
    try:
        result = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None
    return result if result.is_finite() else None


def _json_number(value: Decimal | None) -> float | None:
    return None if value is None else float(value)


def _date_text(value: Any, datemode: int = 0) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and 20_000 <= float(value) <= 80_000:
        if datemode:
            try:
                import xlrd  # type: ignore

                return xlrd.xldate_as_datetime(float(value), datemode).date().isoformat()
            except Exception:
                pass
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date().isoformat()
    text = _text(value)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    match = re.search(r"(20\d{2})[年/.-](\d{1,2})[月/.-](\d{1,2})", text)
    if match:
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3))).isoformat()
    return ""


def _inspection_start_date(value: Any, reference_date: date) -> date | None:
    parsed = _date_text(value)
    if parsed:
        try:
            return date.fromisoformat(parsed)
        except ValueError:
            return None
    text = _text(value)
    match = re.search(r"(?<!\d)(\d{1,2})\s*月\s*(\d{1,2})\s*日?", text)
    if match is None:
        match = re.search(r"(?<!\d)(\d{1,2})[/.](\d{1,2})(?!\d)", text)
    if match is None:
        return None
    try:
        result = date(reference_date.year, int(match.group(1)), int(match.group(2)))
    except ValueError:
        return None
    if result < reference_date - timedelta(days=180):
        try:
            result = result.replace(year=result.year + 1)
        except ValueError:
            return None
    return result


def _sheet_rows(filename: str, content: bytes, *, strict_item_limit: bool = False,
                strict_row_limit: int | None = None) -> list[tuple[str, list[list[Any]], int]]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook  # type: ignore

            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"无法读取 Excel 文件：{exc}") from exc
        result: list[tuple[str, list[list[Any]], int]] = []
        try:
            for sheet in workbook.worksheets:
                if strict_row_limit and sheet.max_row > strict_row_limit:
                    raise HTTPException(422, f"送货工作表“{sheet.title}”超过 {strict_row_limit} 行，请拆分文件后导入")
                if strict_item_limit and "item" in sheet.title.casefold() and sheet.max_row > MAX_IMPORT_ROWS:
                    raise HTTPException(422, f"ITEM 工作表“{sheet.title}”超过 {MAX_IMPORT_ROWS} 行，请拆分或清理空白格式后导入")
                rows = [list(row) for row in sheet.iter_rows(values_only=True, max_row=MAX_IMPORT_ROWS)]
                result.append((sheet.title, rows, 1 if workbook.epoch.year == 1904 else 0))
        finally:
            workbook.close()
        return result
    if suffix == ".xls":
        try:
            import xlrd  # type: ignore

            workbook = xlrd.open_workbook(file_contents=content, formatting_info=False)
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"无法读取旧版 XLS 文件：{exc}") from exc
        if strict_row_limit and any(sheet.nrows > strict_row_limit for sheet in workbook.sheets()):
            raise HTTPException(422, f"送货工作表超过 {strict_row_limit} 行，请拆分文件后导入")
        if strict_item_limit and any("item" in sheet.name.casefold() and sheet.nrows > MAX_IMPORT_ROWS for sheet in workbook.sheets()):
            raise HTTPException(422, f"ITEM 工作表超过 {MAX_IMPORT_ROWS} 行，请拆分后导入")
        return [
            (
                sheet.name,
                [sheet.row_values(index) for index in range(min(sheet.nrows, MAX_IMPORT_ROWS))],
                workbook.datemode,
            )
            for sheet in workbook.sheets()
        ]
    raise HTTPException(status_code=422, detail="该导入类型需要 Excel 文件")


def _header_mapping(rows: list[list[Any]], aliases: dict[str, set[str]]) -> tuple[int, dict[str, int]] | None:
    normalized_aliases = {
        field: {_header_key(alias) for alias in names}
        for field, names in aliases.items()
    }
    best: tuple[int, dict[str, int]] | None = None
    for row_index, row in enumerate(rows[:20]):
        mapping: dict[str, int] = {}
        for column_index, value in enumerate(row):
            key = _header_key(value)
            if not key:
                continue
            for field, names in normalized_aliases.items():
                if field not in mapping and key in names:
                    mapping[field] = column_index
        if best is None or len(mapping) > len(best[1]):
            best = row_index, mapping
    return best


def _cell(row: list[Any], mapping: dict[str, int], field: str) -> Any:
    index = mapping.get(field)
    return row[index] if index is not None and index < len(row) else None


def _split_paper(value: Any) -> tuple[str, str]:
    text = _text(value)
    for packaging_type in PACKAGING_TYPES:
        if text.lower().endswith(packaging_type.lower()):
            quality = text[: -len(packaging_type)].strip(" /-")
            return quality or "待复核", packaging_type
    return text or "待复核", "待复核"


def _specification(row: list[Any], mapping: dict[str, int]) -> str:
    dimensions = [_number(_cell(row, mapping, field)) for field in ("length", "width", "height")]
    present = [str(value.normalize()) for value in dimensions if value is not None]
    return " × ".join(present) + (" in" if present else "")


def _dongkang_destination_factory(value: Any) -> str:
    destination = _header_key(value)
    for suffix in "abcd":
        if destination.startswith(f"华康{suffix}"):
            return f"huakang-{suffix}"
    if destination.startswith("华兴"):
        return "huaxing"
    if destination.startswith("华登"):
        return "huadeng"
    return ""


def _ocr_key(value: Any) -> str:
    return re.sub(r"[^A-Z0-9\u4e00-\u9fff]", "", _text(value).upper())


def _delivery_date_from_text(value: Any) -> str:
    parsed = _date_text(value)
    if parsed:
        return parsed
    text = _text(value)
    match = re.search(r"(?<!\d)(\d{2})[/.\-](\d{1,2})[/.\-](\d{1,2})(?!\d)", text)
    if match is None:
        return ""
    try:
        return date(2000 + int(match.group(1)), int(match.group(2)), int(match.group(3))).isoformat()
    except ValueError:
        return ""


def _delivery_ocr_words(result: Any) -> list[DeliveryOcrWord]:
    if result is None:
        return []
    try:
        texts = list(getattr(result, "txts", ()))
        boxes = list(getattr(result, "boxes", ()))
        scores = list(getattr(result, "scores", ()))
    except Exception:
        return []

    words: list[DeliveryOcrWord] = []
    for index, value in enumerate(texts):
        text = _text(value)
        if not text or index >= len(boxes):
            continue
        try:
            points = boxes[index]
            xs = [float(point[0]) for point in points]
            ys = [float(point[1]) for point in points]
            score = float(scores[index]) if index < len(scores) else 0.8
        except (TypeError, ValueError, IndexError):
            continue
        if not xs or not ys:
            continue
        words.append(
            DeliveryOcrWord(
                text=text,
                left=int(min(xs)),
                top=int(min(ys)),
                right=max(int(min(xs)) + 1, int(max(xs))),
                bottom=max(int(min(ys)) + 1, int(max(ys))),
                confidence=score,
            )
        )
    return words


def _closest_header_word(
    words: list[DeliveryOcrWord],
    predicate,
    header_y: float | None = None,
) -> DeliveryOcrWord | None:
    candidates = [word for word in words if predicate(_ocr_key(word.text))]
    if not candidates:
        return None
    if header_y is None:
        return min(candidates, key=lambda word: word.center_y)
    return min(candidates, key=lambda word: abs(word.center_y - header_y))


def _delivery_table_geometry(words: list[DeliveryOcrWord]) -> dict[str, float] | None:
    order_header = _closest_header_word(
        words,
        lambda key: "订单编号" in key or key in {"PONO", "PON0"},
    )
    description_header = _closest_header_word(words, lambda key: "DESCRIPTION" in key)
    specification_header = _closest_header_word(words, lambda key: "SPECIFICATION" in key)
    if order_header is None or description_header is None or specification_header is None:
        return None

    header_y = median(
        [order_header.center_y, description_header.center_y, specification_header.center_y]
    )
    order_header = _closest_header_word(
        words,
        lambda key: "订单编号" in key or key in {"PONO", "PON0"},
        header_y,
    )
    description_header = _closest_header_word(words, lambda key: "DESCRIPTION" in key, header_y)
    specification_header = _closest_header_word(words, lambda key: "SPECIFICATION" in key, header_y)
    if order_header is None or description_header is None or specification_header is None:
        return None

    quantity_characters = [
        word
        for word in words
        if _ocr_key(word.text) in {"数", "量"} and abs(word.center_y - header_y) <= 90
    ]
    if len(quantity_characters) >= 2:
        quantity_center = sum(word.center_x for word in quantity_characters) / len(quantity_characters)
    else:
        quantity_header = _closest_header_word(words, lambda key: key.startswith("QUANTITY"), header_y)
        if quantity_header is not None:
            quantity_center = quantity_header.left + (quantity_header.right - quantity_header.left) * 0.16
        else:
            quantity_center = specification_header.center_x + (
                specification_header.center_x - description_header.center_x
            ) * 0.78

    unit_price_header = _closest_header_word(words, lambda key: key in {"单价", "UNITPRICE"}, header_y)
    unit_price_center = (
        unit_price_header.center_x
        if unit_price_header is not None
        else quantity_center + (quantity_center - specification_header.center_x) * 0.65
    )
    centers = [
        order_header.center_x,
        description_header.center_x,
        specification_header.center_x,
        quantity_center,
        unit_price_center,
    ]
    if any(right <= left for left, right in zip(centers, centers[1:])):
        return None

    relevant_headers = [
        word
        for word in words
        if abs(word.center_y - header_y) <= 90
        and any(
            marker in _ocr_key(word.text)
            for marker in ("订单编号", "PONO", "DESCRIPTION", "SPECIFICATION", "QUANTITY", "数", "量", "单价")
        )
    ]
    return {
        "header_bottom": float(max(word.bottom for word in relevant_headers)),
        "order_description": (centers[0] + centers[1]) / 2,
        "description_specification": (centers[1] + centers[2]) / 2,
        "specification_quantity": (centers[2] + centers[3]) / 2,
        "quantity_unit_price": (centers[3] + centers[4]) / 2,
    }


def _normalize_order_reference(value: Any) -> str:
    reference = re.sub(r"[^A-Z0-9/\-]", "", _text(value).upper())
    marker = reference.find("700")
    if 0 < marker <= 3 and all(character in "SCOS0528BQD" for character in reference[:marker]):
        reference = "SC" + reference[marker:]
    return reference


def _delivery_order_parts(words: list[DeliveryOcrWord]) -> tuple[str, str, str]:
    candidates = [
        word
        for word in words
        if "/" in word.text and "-" in word.text and re.search(r"\d", word.text)
    ]
    if not candidates:
        return "", "", ""
    primary = max(candidates, key=lambda word: len(_normalize_order_reference(word.text)))
    reference = _normalize_order_reference(primary.text)
    trailing = [
        re.sub(r"\D", "", word.text)
        for word in words
        if word.center_y > primary.center_y
        and word.top <= primary.bottom + max(55, primary.bottom - primary.top)
        and re.fullmatch(r"\s*\d{1,2}\s*", word.text)
    ]
    if trailing:
        reference += trailing[0]
    if "-" not in reference:
        return reference, "", ""
    contract_no, item_no = reference.split("-", 1)
    return reference, contract_no, item_no


def _delivery_description_parts(words: list[DeliveryOcrWord]) -> tuple[str, str]:
    text = " ".join(word.text for word in sorted(words, key=lambda word: word.left)).strip()
    packaging_type = ""
    for candidate in PACKAGING_TYPES:
        if candidate in text:
            packaging_type = candidate
            break
    quality_match = re.search(r"(?<![A-Z0-9])([A-Z]\d{1,3}(?:\+[A-Z0-9]+)?|[A-Z]\d+[A-Z])(?![A-Z0-9])", text.upper())
    paper_quality = quality_match.group(1) if quality_match else ""
    return packaging_type or "待复核", paper_quality or "待复核"


def _delivery_specification(words: list[DeliveryOcrWord], row_words: list[DeliveryOcrWord]) -> str:
    text = " ".join(word.text for word in sorted(words, key=lambda word: word.left))
    match = re.search(
        r"(\d+(?:[.,]\d+)?)\s*[X×*]\s*(\d+(?:[.,]\d+)?)\s*[X×*]\s*(\d+(?:[.,]\d+)?)",
        text,
        re.IGNORECASE,
    )
    if match is None:
        return text.strip()
    dimensions = [part.replace(",", ".") for part in match.groups()]
    has_inches = any(_ocr_key(word.text) in {"IN", "INCH", "INCHES"} for word in row_words)
    return " × ".join(dimensions) + (" in" if has_inches else "")


def _delivery_document_metadata(words: list[DeliveryOcrWord]) -> tuple[str, str]:
    delivery_note_no = ""
    delivery_date = ""
    for word in words:
        if not delivery_note_no:
            match = re.search(r"\bDN\s*[-.:]?\s*([A-Z0-9/-]{5,})", word.text, re.IGNORECASE)
            if match:
                delivery_note_no = f"DN{match.group(1).upper()}"
        if not delivery_date:
            delivery_date = _delivery_date_from_text(word.text)
    return delivery_note_no, delivery_date


def _delivery_rows_from_ocr_words(words: list[DeliveryOcrWord]) -> dict[str, Any] | None:
    geometry = _delivery_table_geometry(words)
    if geometry is None:
        return None

    header_bottom = geometry["header_bottom"]
    total_words = [
        word
        for word in words
        if word.center_y > header_bottom and _ocr_key(word.text) in {"合计", "TOTAL"}
    ]
    table_bottom = min((word.top for word in total_words), default=max(word.bottom for word in words))
    quantity_words = [
        word
        for word in words
        if geometry["specification_quantity"] <= word.center_x < geometry["quantity_unit_price"]
        and header_bottom < word.center_y < table_bottom
        and re.fullmatch(r"\s*\d+(?:[.,]\d+)?\s*", word.text)
        and (_number(word.text) or Decimal(0)) > 0
    ]
    quantity_words.sort(key=lambda word: word.center_y)
    if not quantity_words:
        return None

    grouped_quantities: list[list[DeliveryOcrWord]] = []
    for word in quantity_words:
        if not grouped_quantities:
            grouped_quantities.append([word])
            continue
        current_center = sum(item.center_y for item in grouped_quantities[-1]) / len(grouped_quantities[-1])
        height = max(12, word.bottom - word.top)
        if abs(word.center_y - current_center) <= height * 0.7:
            grouped_quantities[-1].append(word)
        else:
            grouped_quantities.append([word])

    anchors: list[tuple[float, Decimal]] = []
    for group in grouped_quantities:
        raw_amount = "".join(word.text.strip() for word in sorted(group, key=lambda word: word.left))
        amount = _number(raw_amount)
        if amount is not None and amount > 0:
            anchors.append((sum(word.center_y for word in group) / len(group), amount))
    if not anchors:
        return None

    delivery_note_no, delivery_date = _delivery_document_metadata(words)
    rows: list[dict[str, Any]] = []
    for index, (anchor_y, amount) in enumerate(anchors):
        previous_y = anchors[index - 1][0] if index else header_bottom
        next_y = anchors[index + 1][0] if index + 1 < len(anchors) else table_bottom
        top = (previous_y + anchor_y) / 2 if index else header_bottom
        bottom = (anchor_y + next_y) / 2 if index + 1 < len(anchors) else table_bottom
        row_words = [word for word in words if top < word.center_y <= bottom]
        order_words = [word for word in row_words if word.center_x < geometry["order_description"]]
        description_words = [
            word
            for word in row_words
            if geometry["order_description"] <= word.center_x < geometry["description_specification"]
        ]
        specification_words = [
            word
            for word in row_words
            if geometry["description_specification"] <= word.center_x < geometry["specification_quantity"]
        ]
        order_reference, contract_no, item_no = _delivery_order_parts(order_words)
        packaging_type, paper_quality = _delivery_description_parts(description_words)
        rows.append(
            {
                "source_sheet": "OCR四列表格",
                "source_row": index + 1,
                "delivery_note_no": delivery_note_no,
                "delivery_date": delivery_date,
                "order_reference": order_reference,
                "contract_no": contract_no,
                "item_no": item_no,
                "packaging_type": packaging_type,
                "paper_quality": paper_quality,
                "specification": _delivery_specification(specification_words, row_words),
                "delivered_quantity": _json_number(amount),
                "unit_price": 0,
                "location": "",
            }
        )

    complete_orders = sum(1 for row in rows if row["contract_no"] and row["item_no"])
    complete_descriptions = sum(1 for row in rows if row["paper_quality"] != "待复核")
    complete_specifications = sum(1 for row in rows if row["specification"])
    return {
        "rows": rows,
        "delivery_note_no": delivery_note_no,
        "delivery_date": delivery_date,
        "score": len(rows) * 20 + complete_orders * 50 + complete_descriptions * 10 + complete_specifications * 10,
    }


def _parse_delivery_image_table(content: bytes) -> dict[str, Any] | None:
    try:
        from PIL import Image, ImageOps  # type: ignore
        from app.services.carton_mark import get_rapidocr_engine, resize_for_ocr
    except Exception:
        return None

    engine = get_rapidocr_engine()
    if engine is None:
        return None
    try:
        source_image = ImageOps.exif_transpose(Image.open(BytesIO(content))).convert("RGB")
    except Exception:
        return None

    best: dict[str, Any] | None = None
    best_words: list[DeliveryOcrWord] = []
    for angle in (0, 90, 270, 180):
        image = source_image if angle == 0 else source_image.rotate(angle, expand=True, fillcolor="white")
        try:
            result = engine(resize_for_ocr(image))
        except Exception:
            continue
        words = _delivery_ocr_words(result)
        parsed = _delivery_rows_from_ocr_words(words)
        if parsed is not None and (best is None or parsed["score"] > best["score"]):
            best = parsed
            best_words = words
        if best is not None and best["score"] >= 500:
            break

    if best is None or not best["rows"]:
        return None
    missing_orders = sum(1 for row in best["rows"] if not row["contract_no"] or not row["item_no"])
    warnings = [
        "已按送货单四列表格识别：订单编号=合同号+货号、品名=纸品类型+纸质、规格=订单规格、数量=纸箱数量；DN 号仅作为送货单元数据，不参与订单匹配。"
    ]
    if missing_orders:
        warnings.append(f"有 {missing_orders} 行订单编号未完整识别，需要人工复核原图")
    raw_text = "\n".join(word.text for word in sorted(best_words, key=lambda word: (word.center_y, word.left)))
    return {
        "rows": best["rows"],
        "warnings": warnings,
        "engine": "rapidocr-delivery-table-v1",
        "document": {
            "delivery_note_no": best["delivery_note_no"],
            "delivery_date": best["delivery_date"],
            "raw_text_excerpt": raw_text[:3000],
        },
    }


def _item_date_text(value: Any, datemode: int) -> str:
    # ITEM dates must be one complete date; staged/yearless text is human evidence.
    if isinstance(value, str) and not re.fullmatch(
        r"(?:\d{4}[-/.]\d{1,2}[-/.]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{4}|\d{4}年\d{1,2}月\d{1,2}日?)",
        value.strip(),
    ):
        return ""
    try:
        return _date_text(value, datemode)
    except (ValueError, OverflowError):
        return ""


def _parse_item_schedule(sheets: list[tuple[str, list[list[Any]], int]]) -> dict[str, Any]:
    """Read the unified ITEM source without including its duplicate order-summary tab."""
    aliases = {
        "contract_no": {"Contract No.", "合同号"},
        "source_reference": {"SO#/Reference", "SO/Reference"},
        "customer_po": {"P/O#:", "客户PO"},
        "order_type": {"订单类型"},
        "production_no": {"生产单号"},
        "customer_name": {"客名"},
        "item_no": {"产品编号"},
        "product_name": {"产品中文名称"},
        "quantity": {"數量", "数量"},
        "carton_rule": {"装箱"},
        "carton_count": {"箱数"},
        "inspection_window": {"验货日期"},
        "customer_due_date": {"客要求走货期"},
        "note": {"备注"},
    }
    result: list[dict[str, Any]] = []
    mappings: list[dict[str, Any]] = []
    for name, rows, datemode in sheets:
        header = _header_mapping(rows, aliases)
        required = {"contract_no", "order_type", "item_no", "quantity"}
        if header is None or not required.issubset(header[1]):
            raise HTTPException(422, f"ITEM 工作表“{name}”缺少合同号、订单类型、产品编号或数量表头，请检查统一模板")
        index, mapping = header
        header_keys = [_header_key(value) for value in rows[index]]
        for field, names in aliases.items():
            if sum(key in {_header_key(name) for name in names} for key in header_keys) > 1:
                raise HTTPException(422, f"ITEM 工作表“{name}”有重复的 {field} 列，请保留一个明确来源")
        mappings.append({"sheet": name, "header_row": index + 1,
                         "fields": {field: _text(rows[index][column]) for field, column in mapping.items()}})
        for source_row, cells in enumerate(rows[index + 1:], index + 2):
            values = {field: _text(_cell(cells, mapping, field)) for field in mapping}
            if not any(values.get(field) for field in ("contract_no", "source_reference", "item_no", "quantity", "order_type")):
                continue
            quantity = _number(_cell(cells, mapping, "quantity"))
            row: dict[str, Any] = {
                **values, "source_customer_name": values.get("customer_name", ""), "source_sheet": name, "source_row": source_row,
                "template": "unified-item", "quantity": _json_number(quantity),
                "reference": values.get("contract_no", ""),
                "po_numbers": values.get("customer_po", ""),
                "source_inspection_window": values.get("inspection_window", ""),
                "source_customer_due_date": values.get("customer_due_date", ""),
                "inspection_window": _item_date_text(_cell(cells, mapping, "inspection_window"), datemode),
                "customer_due_date": _item_date_text(_cell(cells, mapping, "customer_due_date"), datemode),
            }
            row["date_review_required"] = any(values.get(field) and not row[field]
                                              for field in ("inspection_window", "customer_due_date"))
            if values.get("order_type") != "正单":
                row.update(match_status="REVIEW_REQUIRED", suggestion=f"{values.get('order_type') or '未标订单类型'}：待人工确认，不计入正单漏单核对")
            elif not values.get("contract_no") or not values.get("item_no"):
                row.update(match_status="REVIEW_REQUIRED", suggestion="正单缺少合同号或产品编号，不能按相邻行或单独货号推断订单")
            elif quantity is None or quantity <= 0:
                row.update(match_status="REVIEW_REQUIRED", suggestion="正单数量为空、无效或非正数，请确认需求；不会自动取消采购单")
            result.append(row)
            if len(result) > MAX_IMPORT_ROWS:
                raise HTTPException(422, f"ITEM 数据超过 {MAX_IMPORT_ROWS} 行，请拆分文件后导入")
    if not result:
        raise HTTPException(422, "ITEM 工作表没有可核对的业务明细")
    identities: dict[tuple[str, ...], list[dict[str, Any]]] = {}
    for row in result:
        if row.get("match_status") == "REVIEW_REQUIRED":
            continue
        key = tuple(str(row.get(field) or "").strip().casefold() for field in ("contract_no", "customer_po", "item_no"))
        identities.setdefault(key, []).append(row)
    for duplicates in identities.values():
        if len(duplicates) > 1:
            for row in duplicates:
                row.update(match_status="REVIEW_REQUIRED", suggestion="同合同、客户PO和货号存在多行正单，请先确认是重复还是分批需求；不会自动合并")
    return {"rows": result, "engine": "unified-item-header-mapping", "field_mappings": mappings,
            "warnings": ["仅核对 ITEM 表中的正单；其他订单类型保留待确认。接单表不重复导入。",
                         "Contract No.、SO#/Reference、客户PO分别保留；空白编号不从相邻行补齐。",
                         "无年份、多阶段或无效日期保留原文并标记待确认，不猜测日期。"],
            "review_count": sum(row.get("match_status") == "REVIEW_REQUIRED" or row.get("date_review_required", False) for row in result)}


def _parse_weekly(filename: str, content: bytes) -> dict[str, Any]:
    sheets = _sheet_rows(filename, content, strict_item_limit=True)
    item_sheets = [sheet for sheet in sheets if "item" in sheet[0].casefold()]
    if item_sheets:
        return _parse_item_schedule(item_sheets)
    parsed: list[dict[str, Any]] = []
    warnings: list[str] = []
    for sheet_name, rows, _ in sheets:
        header = _header_mapping(rows, WEEKLY_ALIASES)
        if header is None or not {"contract_no", "quantity"}.issubset(header[1]):
            sheet_key = _header_key(sheet_name)
            if any(marker in sheet_key for marker in ("说明", "示例", "instruction", "example")):
                continue
            if any(any(_text(cell) for cell in row) for row in rows[:20]):
                warnings.append(f"工作表“{sheet_name}”未找到可识别的排期表头")
            continue
        header_index, mapping = header
        for source_row, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            contract_no = _text(_cell(row, mapping, "contract_no"))
            item_no = _text(_cell(row, mapping, "item_no"))
            amount = _number(_cell(row, mapping, "quantity"))
            if not contract_no and not item_no:
                continue
            if amount is None or amount <= 0:
                warnings.append(f"{sheet_name} 第 {source_row} 行数量无效，已跳过")
                continue
            parsed.append(
                {
                    "source_sheet": sheet_name,
                    "source_row": source_row,
                    "reference": contract_no,
                    "contract_no": contract_no,
                    "customer_po": _text(_cell(row, mapping, "customer_po")),
                    "po_numbers": _text(_cell(row, mapping, "po_numbers")),
                    "customer_name": _text(_cell(row, mapping, "customer_name")),
                    "item_no": item_no,
                    "product_name": _text(_cell(row, mapping, "product_name")),
                    "quantity": _json_number(amount),
                    "carton_rule": _text(_cell(row, mapping, "carton_rule")),
                    "inspection_window": _text(_cell(row, mapping, "inspection_window")),
                }
            )
            if len(parsed) >= MAX_PREVIEW_ROWS:
                warnings.append(f"导入预览最多显示 {MAX_PREVIEW_ROWS} 行，其余行未进入本批次")
                return {"rows": parsed, "warnings": warnings, "engine": "excel-header-mapping"}
    if not parsed:
        raise HTTPException(status_code=422, detail="未在文件中找到可识别的每周查货排期明细")
    return {"rows": parsed, "warnings": warnings, "engine": "excel-header-mapping"}


def _parse_delivery_spreadsheet(filename: str, content: bytes, *, strict_rows: bool = False) -> dict[str, Any]:
    parsed: list[dict[str, Any]] = []
    warnings: list[str] = []
    for sheet_name, rows, datemode in _sheet_rows(filename, content,
            strict_row_limit=MAX_PREVIEW_ROWS if strict_rows else None):
        header = _header_mapping(rows, DELIVERY_ALIASES)
        if header is None or not {"item_no", "quantity"}.issubset(header[1]):
            continue
        header_index, mapping = header
        dongkang_format = {_header_key(value) for value in rows[header_index]} >= {"客户单号", "客户料号", "送货单号"}
        for source_row, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            delivery_note_no = _text(_cell(row, mapping, "delivery_note_no"))
            contract_no = _text(_cell(row, mapping, "contract_no"))
            item_no = _text(_cell(row, mapping, "item_no"))
            amount = _number(_cell(row, mapping, "quantity"))
            if not any((delivery_note_no, contract_no, item_no)) and not (strict_rows and any(_text(value) for value in row)):
                continue
            if (amount is None or amount <= 0) and not strict_rows:
                continue
            paper_quality, packaging_type = _split_paper(_cell(row, mapping, "paper_quality"))
            if dongkang_format:
                packaging_type = _text(_cell(row, mapping, "packaging_type")) or packaging_type
            parsed_row = {
                    "source_sheet": sheet_name,
                    "source_row": source_row,
                    "delivery_note_no": delivery_note_no,
                    "delivery_date": _date_text(_cell(row, mapping, "delivery_date"), datemode),
                    "contract_no": contract_no,
                    "customer_po": _text(_cell(row, mapping, "customer_po")),
                    "item_no": item_no,
                    "packaging_type": packaging_type,
                    "paper_quality": paper_quality,
                    "specification": _text(_cell(row, mapping, "specification")) if dongkang_format else _specification(row, mapping),
                    "delivered_quantity": _json_number(amount),
                    "unit_price": _json_number(_number(_cell(row, mapping, "unit_price"))) or 0,
                    "location": _text(_cell(row, mapping, "location")),
                }
            if dongkang_format:
                destination = _text(_cell(row, mapping, "destination"))
                parsed_row.update(template="dongkang-delivery", destination=destination,
                                  destination_factory_id=_dongkang_destination_factory(destination))
                if not item_no:
                    parsed_row.update(match_status="REVIEW_REQUIRED", suggestion="客户料号为空；请对照送货单补齐货号后人工匹配，不按合同号自动关联")
            parsed.append(parsed_row)
            if len(parsed) >= MAX_PREVIEW_ROWS:
                warnings.append(f"导入预览最多显示 {MAX_PREVIEW_ROWS} 行，其余行未进入本批次")
                return {"rows": parsed, "warnings": warnings, "engine": "excel-header-mapping"}
    if not parsed:
        raise HTTPException(status_code=422, detail="未在文件中找到可识别的送货或入库明细")
    return {"rows": parsed, "warnings": warnings, "engine": "excel-header-mapping"}


def _document_text(suffix: str, content: bytes) -> tuple[str, str, str]:
    if suffix == ".pdf":
        from app.services.carton_mark import extract_pdf_text

        text, status = extract_pdf_text(content)
    else:
        from app.services.carton_mark import extract_image_text

        text, status = extract_image_text(content, source="delivery_note")
    return text, status.engine, status.message


def _register_heic_opener() -> None:
    try:
        from pillow_heif import register_heif_opener  # type: ignore

        register_heif_opener(thumbnails=False)
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail="当前后端无法解码 HEIC/HEIF 图片，请安装 pillow-heif 后重试",
        ) from exc


def _parse_delivery_document(filename: str, content: bytes) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix in {".heic", ".heif"}:
        _register_heic_opener()
    if suffix in {".png", ".jpg", ".jpeg", ".heic", ".heif"}:
        table = _parse_delivery_image_table(content)
        if table is not None:
            return table
    text, engine, message = _document_text(suffix, content)
    compact = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not compact:
        raise HTTPException(status_code=422, detail=message or "未能从送货单中识别文字")
    delivery_note_match = re.search(r"\bDN\s*[-.:]?\s*([A-Z0-9/-]{5,})", compact, re.IGNORECASE)
    delivery_note_no = f"DN{delivery_note_match.group(1)}" if delivery_note_match else ""
    date_match = re.search(r"\b(20\d{2})[年/.-](\d{1,2})[月/.-](\d{1,2})", compact)
    delivery_date = (
        date(int(date_match.group(1)), int(date_match.group(2)), int(date_match.group(3))).isoformat()
        if date_match
        else ""
    )
    contract_tokens = re.findall(r"\b(?:SC\d{6,}(?:[/.-]\d+)?|[A-Z]{1,4}\d{5,}(?:-\d+)?)\b", compact, re.IGNORECASE)
    item_tokens = re.findall(r"\b(?:W?[0-9]{4,}[A-Z0-9]*|[A-Z]\d{4,}[A-Z0-9-]*)\b", compact, re.IGNORECASE)
    rows: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for index, contract_no in enumerate(contract_tokens[:50]):
        nearby = item_tokens[index] if index < len(item_tokens) else ""
        key = (_identity(contract_no), _identity(nearby))
        if key in seen:
            continue
        seen.add(key)
        rows.append(
            {
                "source_sheet": "OCR",
                "source_row": index + 1,
                "delivery_note_no": delivery_note_no,
                "delivery_date": delivery_date,
                "contract_no": contract_no,
                "item_no": nearby,
                "packaging_type": "待复核",
                "paper_quality": "待复核",
                "specification": "",
                "delivered_quantity": 0,
                "unit_price": 0,
                "location": "",
            }
        )
    warnings = ["图片/PDF 仅作为 OCR 预览，数量和纸品字段必须逐行人工复核"]
    if message:
        warnings.append(message)
    return {
        "rows": rows,
        "warnings": warnings,
        "engine": engine,
        "document": {
            "delivery_note_no": delivery_note_no,
            "delivery_date": delivery_date,
            "raw_text_excerpt": compact[:3000],
        },
    }


def _match_rows(db: Session, factory_id: str, import_type: str, rows: list[dict[str, Any]]) -> None:
    joined = list(
        db.execute(
            select(CartonOrderLine, CartonOrder)
            .join(CartonOrder, CartonOrder.id == CartonOrderLine.order_id)
            .where(
                CartonOrder.factory_id == factory_id,
                CartonOrder.status != "CANCELLED",
            )
        ).all()
    )
    all_joined = joined
    for row in rows:
        if row.get("match_status") == "REVIEW_REQUIRED":
            continue
        if row.get("template") == "dongkang-delivery":
            destination_factory = row.get("destination_factory_id")
            if not destination_factory or destination_factory != factory_id:
                row.update(match_status="REVIEW_REQUIRED", suggestion=(
                    f"送货对象“{row.get('destination') or '空'}”与当前厂区不符或无法识别；请在正确厂区导入并人工核对"))
                continue
        po_key = str(row.get("customer_po") or "").strip().casefold()
        joined = [(line, order) for line, order in all_joined if not po_key or order.customer_po.strip().casefold() == po_key]
        if row.get("template") == "unified-item":
            identity = lambda value: _text(value).casefold()
        elif row.get("template") == "dongkang-delivery":
            identity = _dongkang_identity
        else:
            identity = _identity
        contract_key = identity(row.get("contract_no"))
        item_key = identity(row.get("item_no"))
        exact = [
            (line, order)
            for line, order in joined
            if contract_key
            and item_key
            and identity(order.contract_no) == contract_key
            and identity(order.item_no) == item_key
        ]
        candidates = exact
        basis = "合同号 + 货号" if exact else ""
        if not candidates and contract_key and row.get("template") not in {"unified-item", "dongkang-delivery"}:
            by_contract = [(line, order) for line, order in joined if _identity(order.contract_no) == contract_key]
            if len({order.id for _, order in by_contract}) == 1:
                candidates = by_contract
                basis = "合同号"
        if not candidates and item_key and row.get("template") not in {"unified-item", "dongkang-delivery"}:
            by_item = [(line, order) for line, order in joined if _identity(order.item_no) == item_key]
            if len({order.id for _, order in by_item}) == 1:
                candidates = by_item
                basis = "唯一货号"

        # Missing PO must not be silently resolved by material differences across orders.
        ambiguous_po_orders = not po_key and len({order.id for _, order in exact}) > 1 and any(order.customer_po for _, order in exact)
        if ambiguous_po_orders:
            row["match_status"] = "AMBIGUOUS"
            row["suggestion"] = "同合同货号存在多个客户 PO，请填写客户 PO 或人工选择订单"
            continue
        if import_type == "DELIVERY_NOTE" and candidates:
            packaging = _text(row.get("packaging_type"))
            paper_quality = _text(row.get("paper_quality"))
            specification = _dimension_identity(row.get("specification"))
            narrowing_basis: list[str] = []
            if paper_quality not in {"", "待复核"}:
                narrowed = [
                    pair for pair in candidates
                    if pair[0].paper_quality.lower() == paper_quality.lower()
                ]
                if narrowed:
                    candidates = narrowed
                    narrowing_basis.append("纸质")
            if specification:
                narrowed = [
                    pair for pair in candidates
                    if _dimension_identity(pair[0].specification) == specification
                ]
                if narrowed:
                    candidates = narrowed
                    narrowing_basis.append("规格")
            if packaging not in {"", "待复核", "普通箱", "纸箱"}:
                narrowed = [pair for pair in candidates if pair[0].packaging_type == packaging]
                if narrowed:
                    candidates = narrowed
                    narrowing_basis.append("纸品类型")
            if narrowing_basis:
                basis += " + " + " + ".join(narrowing_basis)

        unique_orders = {order.id: order for _, order in candidates}
        if import_type in {"WEEKLY_SCHEDULE", "INSPECTION_SCHEDULE"}:
            if len(unique_orders) == 1:
                order = next(iter(unique_orders.values()))
                row.update(
                    {
                        "order_id": order.id,
                        "order_no": order.order_no,
                        "customer_code": order.customer_code,
                        "customer_name": order.customer_name,
                        "order_status": order.status,
                        "match_basis": basis,
                    }
                )
                schedule_quantity = _number(row.get("quantity")) or Decimal(0)
                if import_type == "WEEKLY_SCHEDULE" and (order.product_order_quantity is None or schedule_quantity != Decimal(order.product_order_quantity)):
                    row["match_status"] = "QUANTITY_MISMATCH"
                    row["suggestion"] = f"排期数量 {schedule_quantity} 与订单数量 {order.product_order_quantity} 不一致，请人工确认"
                else:
                    row["match_status"] = "MATCHED"
                    row["suggestion"] = "已匹配正式纸箱订单" if import_type == "WEEKLY_SCHEDULE" else "已关联正式纸箱订单，等待计算交货提醒"
                if row.get("template") == "unified-item":
                    row["procurement_state"] = (
                        "NEEDS_ORDER" if order.status in {"DRAFT", "CONFIRMED"}
                        else "REVIEW" if order.product_order_quantity is None
                        else "COMPLETED" if order.status == "COMPLETED" else "ORDERED"
                    )
                    if order.product_order_quantity is not None and schedule_quantity > Decimal(order.product_order_quantity):
                        row["procurement_state"] = "NEEDS_ORDER"
                    source_due = str(row.get("customer_due_date") or "")
                    if source_due and order.customer_due_date and source_due != order.customer_due_date:
                        row["date_difference"] = {"business_date": source_due, "order_date": order.customer_due_date}
                        row["suggestion"] += f"；业务走货期 {source_due} 与订单客户交期 {order.customer_due_date} 不一致，请确认"
                        if row["match_status"] == "MATCHED":
                            row["match_status"] = "DATE_MISMATCH"
            else:
                row["match_status"] = "AMBIGUOUS" if candidates else "MISSING_ORDER"
                if row.get("template") == "unified-item":
                    row["procurement_state"] = "REVIEW" if candidates else "NEEDS_ORDER"
                row["suggestion"] = "找到多个候选订单，请人工选择" if candidates else (
                    "查货合同未找到正式纸箱订单，请先核对是否漏单"
                    if import_type == "INSPECTION_SCHEDULE"
                    else "未找到正式纸箱订单，仅生成异常待办"
                )
            if row.get("date_review_required"):
                row["suggestion"] += "；日期原文需人工确认，不自动计算交期"
                if row["match_status"] == "MATCHED":
                    row["match_status"] = "REVIEW_REQUIRED"
            continue

        if len(candidates) == 1:
            line, order = candidates[0]
            row.update(
                {
                    "order_line_id": line.id,
                    "order_id": order.id,
                    "order_no": order.order_no,
                    "customer_code": order.customer_code,
                    "customer_name": order.customer_name,
                    "contract_no": order.contract_no,
                    "customer_po": order.customer_po,
                    "item_no": order.item_no,
                    "packaging_type": line.packaging_type,
                    "paper_quality": line.paper_quality,
                    "specification": line.specification,
                    "unit": line.unit,
                    "match_basis": basis,
                    "match_status": "MATCHED",
                }
            )
            if row.get("template") == "dongkang-delivery":
                row["order_unit_price"] = _json_number(_number(line.unit_price))
                supplier_price = _number(row.get("unit_price"))
                order_price = _number(line.unit_price)
                if supplier_price is not None and order_price is not None and supplier_price != order_price:
                    row["suggestion"] = f"送货单单价 {supplier_price} 与采购订单价 {order_price} 不同；入库默认采用采购订单价，请核对"
        else:
            row["match_status"] = "AMBIGUOUS" if candidates else "MISSING_ORDER"
            if candidates:
                row["suggestion"] = "找到多条纸品明细，请按纸质和规格人工选择"
            elif not contract_key or not item_key:
                row["suggestion"] = "订单编号（合同号+货号）未完整识别，请对照原送货单人工复核"
            else:
                row["suggestion"] = "未找到可关联的正式订单明细"


def _apply_inspection_reminders(
    rows: list[dict[str, Any]],
    *,
    advance_days: int,
    reference_date: date,
) -> None:
    for row in rows:
        if row.get("match_status") == "REVIEW_REQUIRED" or row.get("date_review_required"):
            row["reminder_status"] = "REVIEW_REQUIRED"
            continue
        inspection_date = _inspection_start_date(row.get("inspection_window"), reference_date)
        row["advance_days"] = advance_days
        if inspection_date is None:
            row.update(
                {
                    "inspection_start_date": "",
                    "required_delivery_date": "",
                    "days_until_delivery": None,
                    "reminder_status": "INVALID_DATE",
                    "suggestion": "无法识别验货开始日期，请人工补充后再计算最迟交货日",
                }
            )
            continue

        required_date = inspection_date - timedelta(days=advance_days)
        days_until = (required_date - reference_date).days
        row.update(
            {
                "inspection_start_date": inspection_date.isoformat(),
                "required_delivery_date": required_date.isoformat(),
                "days_until_delivery": days_until,
            }
        )
        if row.get("match_status") == "MISSING_ORDER":
            row["reminder_status"] = "MISSING_ORDER"
            row["suggestion"] = f"最迟应于 {required_date.isoformat()} 前到纸箱；当前未找到正式订单，请立即核对漏单"
        elif row.get("match_status") == "AMBIGUOUS":
            row["reminder_status"] = "AMBIGUOUS"
            row["suggestion"] = f"最迟应于 {required_date.isoformat()} 前到纸箱；存在多个候选订单，请人工关联"
        elif row.get("order_status") == "COMPLETED":
            row["reminder_status"] = "READY"
            row["suggestion"] = "纸箱订单已全部收齐，无需催交"
        elif days_until < 0:
            row["reminder_status"] = "OVERDUE"
            row["suggestion"] = f"最迟交货日 {required_date.isoformat()} 已逾期 {abs(days_until)} 天，请立即跟进纸箱到料"
        elif days_until <= 1:
            row["reminder_status"] = "DUE_SOON"
            row["suggestion"] = f"最迟交货日为 {required_date.isoformat()}，请优先确认纸箱到料安排"
        else:
            row["reminder_status"] = "UPCOMING"
            row["suggestion"] = f"请在 {required_date.isoformat()} 前完成纸箱交货，距最迟交货日 {days_until} 天"


def parse_carton_import(
    db: Session,
    factory_id: str,
    import_type: str,
    filename: str,
    content: bytes,
    options: dict[str, Any] | None = None,
) -> dict[str, Any]:
    options = options or {}
    suffix = Path(filename).suffix.lower()
    if import_type in {"WEEKLY_SCHEDULE", "INSPECTION_SCHEDULE"}:
        parsed = _parse_weekly(filename, content)
    elif suffix in {".xlsx", ".xlsm", ".xls"}:
        parsed = _parse_delivery_spreadsheet(filename, content)
    else:
        parsed = _parse_delivery_document(filename, content)
    rows = parsed.get("rows", [])
    _match_rows(db, factory_id, import_type, rows)
    if import_type == "INSPECTION_SCHEDULE":
        advance_days = max(0, min(30, int(options.get("advance_days", 3))))
        try:
            reference_date = date.fromisoformat(str(options.get("reference_date") or ""))
        except ValueError:
            reference_date = date.today()
        _apply_inspection_reminders(
            rows,
            advance_days=advance_days,
            reference_date=reference_date,
        )
    matched = sum(1 for row in rows if row.get("match_status") == "MATCHED")
    issues = len(rows) - matched
    reminder_rows = [row for row in rows if row.get("reminder_status") not in {None, "READY"}]
    parsed.update(
        {
            "message": (
                "下周查货合同已解析并按最迟交货日生成提醒；不会自动创建订单或库存"
                if import_type == "INSPECTION_SCHEDULE"
                else "导入内容已解析并生成待复核预览；不会自动创建订单或库存"
            ),
            "creates_order": False,
            "creates_inventory": False,
            "row_count": len(rows),
            "matched_count": matched,
            "issue_count": issues,
            "reminder_count": len(reminder_rows),
            "overdue_count": sum(1 for row in rows if row.get("reminder_status") == "OVERDUE"),
            "due_soon_count": sum(1 for row in rows if row.get("reminder_status") == "DUE_SOON"),
            "ready_count": sum(1 for row in rows if row.get("reminder_status") == "READY"),
            "advance_days": int(options.get("advance_days", 3)) if import_type == "INSPECTION_SCHEDULE" else None,
            "parser_version": DELIVERY_IMPORT_PARSER_VERSION if import_type == "DELIVERY_NOTE" else SCHEDULE_IMPORT_PARSER_VERSION,
        }
    )
    return parsed
