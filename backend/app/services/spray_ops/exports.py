"""Bounded exports built outside write locks, then frozen under revision control."""
from io import BytesIO
from decimal import Decimal
from datetime import datetime, timezone, timedelta
import hashlib

from sqlalchemy import select
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

from app.models import spray_ops as m
from .common import add, get, scoped, serialize, require, digest, project
from .production import rows
from . import finance

TITLES = {"plan": "喷油部生产计划表", "daily": "生产日报表", "stock": "胶件库存明细", "purchase": "油漆采购单", "material": "油漆批次进耗存", "delivery": "喷油部送货明细", "settlement": "喷油内部月结明细", "request": "喷油内部请款表", "operating": "喷油部经营收支核对"}
FACTORY_NAMES = dict(huaxing="华兴", huadeng="华登", **{"huakang-a": "华康A", "huakang-b": "华康B"})
STATE_NAMES = dict(planned="已排待开", started="执行中", paused="暂停", completed="已完成", cancelled="已取消", white="白件", wip="在制", finished="合格成品", hold="质量待判", complete="已登记", partial="部分登记", unknown="待核对")


def display_time(value):
    return datetime.fromisoformat(value).astimezone(timezone(timedelta(hours=8))).strftime('%Y-%m-%d %H:%M') if value else ""


def permissions(kind):
    if kind in {"settlement", "request"}:
        return ("export", "settle", "cost_read")
    if kind == "operating":
        return ("export", "cost_read", "payroll_read")
    if kind in {"purchase", "material"}:
        return ("export", "cost_read")
    return ("export",)


def prepare(db, f, b, *, cost, payroll):
    """No pagination truncation: filters apply to all selected facts."""
    def filtered(model):
        stmt = scoped(db, model, f)
        if hasattr(model, "business_date"):
            if b.from_date:
                stmt = stmt.where(model.business_date >= str(b.from_date))
            if b.to_date:
                stmt = stmt.where(model.business_date <= str(b.to_date))
        if b.entity_id:
            stmt = stmt.where(model.id == b.entity_id)
        data = list(db.scalars(stmt.order_by(model.created_at, model.id).limit(100001)))
        require(len(data) <= 100000, "export_size", "请缩小筛选范围至十万条记录以内", 413)
        return data

    output, headers, subtitle, daily_payroll = [], [], "", []
    if b.kind == "daily":
        headers = ["生产日期", "班次", "资源", "货号 / 部件", "工序", "目标数", "正班生产数", "加班生产数", "合格", "待判", "报废", "单位", "姓名", "正班工时", "加班工时", "非生产工时", "个人产量"]
        if payroll:
            headers += ["已确认生产计薪", "计薪币种"]
        headers += ["备注", "签名"]
        tasks = {task.id: task for task in rows(db, m.SprayOpsTask, f)}
        resources = {r.id: r for r in rows(db, m.SprayOpsResource, f)}
        steps = {step.id: step for step in rows(db, m.SprayOpsStep, f)}
        demand_lines = {line.id: line for line in rows(db, m.SprayOpsDemandLine, f)}
        employees = {e.id: e for e in rows(db, m.SprayOpsEmployee, f)}
        official = {p.id: p for p in rows(db, m.SprayOpsPayroll, f, status="confirmed")} if payroll else {}
        pay_lines = {p.id: p for p in rows(db, m.SprayOpsPayrollLine, f) if p.payroll_id in official} if payroll else {}
        pay_allocations = {a.labor_id: (a.payroll_amount, official[pay_lines[a.payroll_line_id].payroll_id].currency) for a in rows(db, m.SprayOpsPayrollAllocation, f) if a.payroll_line_id in pay_lines} if payroll else {}
        daily_reports = [report for report in filtered(m.SprayOpsReport) if report.status == "confirmed"]
        if payroll:
            days = {report.business_date for report in daily_reports}
            for line in pay_lines.values():
                parent = official[line.payroll_id]
                if parent.business_date in days:
                    daily_payroll.append([parent.business_date, employees[line.employee_id].name, line.base_amount, line.guarantee, line.subsidy, line.adjustment, line.payroll_amount, parent.currency, line.evidence, ""])
        for report in daily_reports:
            if report.status != "confirmed":
                continue
            for row in rows(db, m.SprayOpsReportRow, f, report_id=report.id):
                task = tasks[row.task_id]
                labor = rows(db, m.SprayOpsLabor, f, row_id=row.id, active=True)
                identities = [demand_lines[a.line_id] for a in rows(db, m.SprayOpsTaskAllocation, f, task_id=task.id)]
                for index, person in enumerate(labor or [None]):
                    values = [report.business_date, report.shift, resources[task.resource_id].name, " / ".join(dict.fromkeys(line.item_no + " " + line.part for line in identities)), steps[task.step_id].name]
                    # A team produces this quantity once. Extra employee lines
                    # carry their own hours, never cloned production quantities.
                    values += [task.quantity, row.normal, row.overtime, row.good, row.hold, row.scrap] if index == 0 else [""] * 6
                    values += [steps[task.step_id].input_unit, employees[person.employee_id].name if person else "员工待核对", person.hours if person else None, person.overtime_hours if person else None, person.nonproductive_hours if person else None, person.personal_quantity if person else None]
                    if payroll:
                        values += list(pay_allocations.get(person.id if person else "", (None, "未正式计薪")))
                    values += [row.reason + ("；班组产量仅首行列示" if len(labor) > 1 else ""), ""]
                    output.append(values)
    elif b.kind == "plan":
        headers = ["资源", "货号", "部件", "订单", "工序", "计划数量", "单位", "开始时间", "结束时间", "实绩", "状态", "备注"]
        for task in filtered(m.SprayOpsTask):
            if task.status == "cancelled":
                continue
            if b.from_date and task.start_at[:10] < str(b.from_date) or b.to_date and task.start_at[:10] > str(b.to_date):
                continue
            resource, step = get(db, m.SprayOpsResource, f, task.resource_id), get(db, m.SprayOpsStep, f, task.step_id)
            for allocation in rows(db, m.SprayOpsTaskAllocation, f, task_id=task.id):
                line = get(db, m.SprayOpsDemandLine, f, allocation.line_id)
                demand = get(db, m.SprayOpsDemand, f, line.demand_id)
                output.append([resource.name, line.item_no, line.part, demand.document_no, step.name, allocation.quantity, step.input_unit, display_time(task.start_at), display_time(task.end_at), allocation.consumed, STATE_NAMES.get(task.status, task.status), task.note])
    elif b.kind in {"settlement", "request"}:
        require(b.entity_id, "settlement_required", "请指定已冻结的月结单", 422)
        settlement = get(db, m.SprayOpsSettlement, f, b.entity_id)
        subtitle = f"{settlement.month} · {settlement.document_no} · {settlement.counterparty} · {settlement.currency}"
        headers = ["月结单号", "月份", "往来方", "货号 / 部件", "送货单号", "本次结算数量", "单价", "金额", "币种"]
        for line in rows(db, m.SprayOpsSettlementLine, f, settlement_id=settlement.id):
            delivery_line = get(db, m.SprayOpsDeliveryLine, f, line.delivery_line_id)
            delivery = get(db, m.SprayOpsDelivery, f, delivery_line.delivery_id)
            stock = get(db, m.SprayOpsStock, f, delivery_line.stock_id)
            demand_line = get(db, m.SprayOpsDemandLine, f, stock.line_id)
            output.append([settlement.document_no, settlement.month, settlement.counterparty, demand_line.item_no + " " + demand_line.part, delivery.document_no, line.quantity, line.price, line.amount, settlement.currency])
    elif b.kind == "stock":
        headers = ["批次", "订单", "货号", "部件", "状态", "数量", "预留数量", "单位", "可转序时间"]
        for stock in filtered(m.SprayOpsStock):
            batch = get(db, m.SprayOpsBatch, f, stock.batch_id)
            line = get(db, m.SprayOpsDemandLine, f, stock.line_id) if stock.line_id else None
            order = get(db, m.SprayOpsDemand, f, line.demand_id) if line else None
            output.append([batch.document_no, order.document_no if order else "未匹配", batch.item_no, batch.part, STATE_NAMES.get(stock.state, stock.state), stock.quantity, stock.reserved, stock.unit, display_time(stock.ready_at)])
    elif b.kind == "purchase":
        headers = ["采购单号", "供应商", "日期", "物料", "数量", "单位", "单价", "金额", "币种", "交货期", "已收", "税口径"]
        for purchase in filtered(m.SprayOpsPurchase):
            for line in rows(db, m.SprayOpsPurchaseLine, f, purchase_id=purchase.id):
                material = get(db, m.SprayOpsMaterial, f, line.material_id)
                output.append([purchase.document_no, purchase.supplier, purchase.business_date, material.code + " " + material.name, line.quantity, line.unit, line.price, line.quantity * line.price if line.price is not None else "未定价", purchase.currency, line.due_date, line.received, purchase.tax_basis])
    elif b.kind == "material":
        headers = ["入库单号", "日期", "物料", "实收", "仓库", "车间", "耗用", "单位", "批次单位成本", "耗用成本", "币种"]
        for lot in filtered(m.SprayOpsMaterialLot):
            receipt = get(db, m.SprayOpsMaterialReceipt, f, lot.receipt_id)
            if b.from_date and receipt.business_date < str(b.from_date) or b.to_date and receipt.business_date > str(b.to_date):
                continue
            material = get(db, m.SprayOpsMaterial, f, lot.material_id)
            output.append([receipt.document_no, receipt.business_date, material.code + " " + material.name, lot.quantity, lot.warehouse, lot.floor, lot.consumed, lot.unit, lot.unit_cost, lot.unit_cost * lot.consumed if lot.unit_cost is not None else "成本待确认", lot.currency])
    elif b.kind == "delivery":
        headers = ["送货单号", "日期", "往来方", "货号 / 部件", "交货数", "验收", "拒收", "退回", "在途", "单位", "备注"]
        for delivery in filtered(m.SprayOpsDelivery):
            for line in rows(db, m.SprayOpsDeliveryLine, f, delivery_id=delivery.id):
                stock = get(db, m.SprayOpsStock, f, line.stock_id)
                batch = get(db, m.SprayOpsBatch, f, stock.batch_id)
                output.append([delivery.document_no, delivery.business_date, delivery.counterparty, batch.item_no + " " + batch.part, line.quantity, line.accepted, line.rejected, line.returned, line.quantity - line.accepted - line.rejected, stock.unit, delivery.warehouse_ref])
    else:
        require(b.month, "month_required", "经营报表须选择月份", 422)
        summary = finance.economics(db, f, b.month, b.currency)
        subtitle = b.month + " · " + b.currency
        headers = ["月份", "币种", "项目", "金额", "核对状态"]
        for key, label in [("operating_value", "工序产值"), ("payroll_amount", "员工工资"), ("material_cost", "实际材料耗用"), ("expense_amount", "费用净额"), ("surplus", "管理口径结余")]:
            output.append([b.month, b.currency, label, summary[key], "待核对" if summary[key] is None else "已登记金额；完整性见覆盖度"])
        for key, state in summary["coverage"].items():
            output.append([b.month, b.currency, {"production":"生产资料覆盖", "pricing":"产值定价覆盖", "payroll":"工资确认覆盖", "material_cost":"材料成本覆盖"}.get(key, key), None, STATE_NAMES.get(state, state)])
    required = list(permissions(b.kind))
    if b.kind == "daily" and payroll:
        required.append("payroll_read")
    return dict(kind=b.kind, title=TITLES[b.kind], factory=FACTORY_NAMES[f], headers=headers, rows=output, subtitle=subtitle, daily_payroll=daily_payroll, from_date=str(b.from_date or ""), to_date=str(b.to_date or ""), required_permissions=required, generated_at=datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC'), base_revision=b.base_revision)


def render(snapshot):
    workbook = Workbook()
    workbook.remove(workbook.active)
    if snapshot['kind'] == 'daily':
        production, labor = [], []
        group = None
        for row in snapshot['rows']:
            if row[5] != '':
                group = [*row[:12], 0, row[-2]]
                production.append(group)
            if group is not None and row[12] != '员工待核对':
                group[12] += 1
            labor.append([*row[:5], *row[12:-2], row[-1]])
        render_sheet(workbook, {**snapshot, 'headers': [*snapshot['headers'][:12], '人数', '备注'], 'rows': production, 'widths': [11,8,11,15,12,7,8,8,7,7,7,5,5,18]})
        render_sheet(workbook, {**snapshot, 'title': '员工工时与计薪', 'headers': [*snapshot['headers'][:5], *snapshot['headers'][12:-2], '签名'], 'rows': labor, 'widths': [11,8,11,15,12,10,8,8,8,8,12,7,10]})
        if 'payroll_read' in snapshot['required_permissions']:
            render_sheet(workbook, {**snapshot, 'title': '当日工资核对', 'headers': ['生产日期','姓名','基础工资','补足保底','补贴','调整','当日合计','币种','调整依据','签名'], 'rows': snapshot['daily_payroll'], 'subtitle':'按完整生产日汇总，保底与补贴每天只计一次', 'widths':[12,12,11,10,10,10,12,8,25,12]})
    else:
        render_sheet(workbook, snapshot)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_sheet(workbook, snapshot):
    sheet = workbook.create_sheet()
    sheet.title = snapshot["title"][:31]
    columns = len(snapshot["headers"])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    sheet.cell(1, 1, snapshot["title"])
    sheet.cell(1, 1).font = Font(name="Microsoft YaHei", size=18, bold=True, color="164D46")
    sheet.row_dimensions[1].height = 34
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    sheet.cell(2, 1, snapshot["factory"] + " · " + (snapshot.get("subtitle") or (snapshot["from_date"] or "全部") + " 至 " + (snapshot["to_date"] or "全部")))
    sheet.append(snapshot["headers"])
    for values in snapshot["rows"]:
        # All user-origin strings are text, never executable Excel formula cells.
        sheet.append([("'" + value if value.startswith(("=", "+", "-", "@")) else value) if isinstance(value, str) else value if value is not None else "待确认" for value in values])
    thin = Side(style="thin", color="D4E1DE")
    for row in sheet.iter_rows(min_row=3):
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color="FFFFFF" if cell.row == 3 else "243E39", bold=cell.row == 3)
            cell.fill = PatternFill("solid", fgColor="17685C" if cell.row == 3 else "F1F7F5" if cell.row % 2 == 0 else "FFFFFF")
            cell.border = Border(bottom=thin)
            numeric = cell.row > 3 and cell.data_type == "n"
            header = snapshot["headers"][cell.column - 1]
            money = any(word in header for word in ("金额", "单价", "成本", "工资", "计薪", "保底", "补贴", "调整", "合计"))
            cell.alignment = Alignment(vertical="center", horizontal="right" if numeric else "center" if header in {"单位", "币种", "计薪币种"} else "left", wrap_text=not numeric)
            if numeric:
                places = min(6, max(2 if money else 0, -Decimal(str(cell.value)).normalize().as_tuple().exponent))
                cell.number_format = "#,##0" + ("." + "0" * places if places else "")
            elif cell.data_type == "s":
                cell.number_format = "@"
        sheet.row_dimensions[row[0].row].height = 34 if snapshot["kind"] == "daily" else 29
    from openpyxl.utils import get_column_letter
    for col in range(1, columns + 1):
        header = snapshot["headers"][col - 1]
        sheet.column_dimensions[get_column_letter(col)].width = snapshot['widths'][col-1] if 'widths' in snapshot else 24 if any(word in header for word in ["备注", "货号", "物料", "时间", "核对"]) else 16
    footer = sheet.max_row + 2
    sheet.cell(footer, 1, "填表人：")
    sheet.cell(footer, max(2, columns // 2), "审核人：")
    sheet.cell(footer, max(3, columns - 1), "签收 / 核对：")
    sheet.merge_cells(start_row=footer+2, start_column=1, end_row=footer+2, end_column=columns)
    sheet.cell(footer+2, 1, f"生成：{snapshot.get('generated_at', '')} · 数据修订：{snapshot.get('base_revision', '')} · 冻结快照；业务时间为北京时间，缺项与执行状态见明细")
    sheet.cell(footer+2, 1).font = Font(name="Microsoft YaHei", size=8, color="61736E")
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:{get_column_letter(columns)}{max(3, footer-2)}"
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:3"
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins = PageMargins(left=.25, right=.25, top=.35, bottom=.35, header=.15, footer=.15)
    sheet.print_area = f"A1:{get_column_letter(columns)}{footer+3}"
    sheet.oddFooter.center.text = "第 &P 页 / 共 &N 页"


def freeze(db, f, b, snapshot, payload, gate):
    require(b.expected_version == 0, "version_conflict", "新导出版本应为 0")
    require(gate.revision == b.base_revision, "stale_export", "导出准备期间数据已变化，请刷新后重试")
    artifact = add(db, m.SprayOpsExport, f, name=f"{snapshot['factory']}-{snapshot['title']}.xlsx", kind=b.kind, payload=payload, sha256=hashlib.sha256(payload).hexdigest(), source_fingerprint=digest(snapshot), required_permissions=snapshot['required_permissions'])
    return {key: value for key, value in serialize(artifact).items() if key != "payload"}
