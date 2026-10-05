"""Bounded, read-only previews and atomic factory-scoped master XLSX imports."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
import hashlib
import json
import re
import unicodedata
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
from sqlalchemy import select

from app.models.carton_master import CartonMasterRecord as Record
from app.models.carton_positions import CartonLocation
from app.schemas.carton_master import MasterSave
from app.services import carton_master as master

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000
HEADERS = {
    "paper-options": ["纸品类型", "纸质", "长", "宽", "高"],
    "configurations": ["货号", "产品名称", "配置组", "纸品类型", "纸质", "长", "宽", "高", "尺寸单位", "计量单位", "每箱装产品数", "推荐", "状态", "备注"],
    "locations": ["仓库", "仓位或范围"],
}


def template(kind):
    book = Workbook()
    sheet = book.active
    sheet.title = "导入数据"
    sheet.append(HEADERS[kind])
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 20
        sheet.column_dimensions[column[0].column_letter].number_format = "@"
        column[0].font = Font(bold=True)
    notes = book.create_sheet("说明与示例")
    messages = ["仅导入“导入数据”页；本页示例不会导入。支持 xlsx / xlsm，最多 5 MB、1000 行。",
                    "纸品类型与纸质每格仅填一项，不使用逗号、顿号、分号或换行拼接。禁止公式，请粘贴为值。",
                    "长、宽、高为正数，最多 8 位小数；纸品选项允许三项全空，货号包装必须全部填写。",
                    "空行忽略；货号、产品名称、单位、装箱数须显式填写，不从相邻行推测。货号必须是真实文本，保留前导零；数字货号一律拒绝，仅修改显示格式无效，请先设为文本后重新输入或以单引号开头输入。",
                    "相同货号+配置组组成一套多纸品配置；配置组空白统一为“默认”。同组产品、推荐、状态和备注必须一致。",
                    "推荐填 是/否（空白为否），状态填 启用/停用（空白为启用）；同货号最多一个启用推荐组。",
                    "不同配置追加为包装变体；完全相同配置跳过，保留现有状态、推荐和备注。客户与价格不导入。"]
    if kind == "locations":
        messages = ["仅导入“导入数据”页；本页示例不会导入。支持 xlsx / xlsm，最多 5 MB、1000 数据行；所有行展开后合计最多 1000 个仓位（含跳过项）。",
                    "两列均须填写真实文本，禁止公式和数字单元格；请先设为文本后重新输入，或以单引号开头输入，保留 001 等前导零。",
                    "仓库每格仅一个名称；仓位可填单项或列表，列表用顿号、逗号、分号或换行分隔，如 A1-A25、B01-B03。",
                    "单个仓位支持中英文、数字、下划线及 A-00 形式，不含空白。范围须同前缀且升序，两端均写完整，如 A1-A25 或 A-01-A-03。",
                    "带前导零时范围两端数字宽度必须一致：B01-B03 保留两位；A1-A03、A01-A3 均拒绝。",
                    "仓库和仓位自动去除首尾空白并转大写；文件内重复报错。已有仓位（包括停用）仅跳过，不改状态、库存或历史。",
                    "新增仓位为启用状态；首个仓位同时建立对应仓库，不创建空仓。禁止导入系统待核仓位或 CL-UNKNOWN 标识。"]
    for message in messages:
        notes.append([message])
        row = notes.max_row
        notes.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(HEADERS[kind]))
        notes.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    notes.append(HEADERS[kind])
    header_row = notes.max_row
    for cell in notes[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if kind == "locations":
        notes.append(["A车间", "A1-A25、B01-B03"])
    elif kind == "paper-options":
        notes.append(["外箱", "A33", 30, 20, 15])
    else:
        notes.append(["00123", "示例产品", "标准装", "外箱", "A33", 30, 20, 15, "cm", "个", 24, "是", "启用", "同一货号多纸品示例"])
        notes.append(["00123", "示例产品", "标准装", "内箱", "B25", 18, 12.5, 10, "cm", "个", 24, "是", "启用", "同一货号多纸品示例"])
    widths = ([18, 22, 14, 14, 14, 10, 10, 10, 12, 12, 16, 10, 10, 24]
              if kind == "configurations" else [20] * len(HEADERS[kind]))
    for index, width in enumerate(widths, 1):
        notes.column_dimensions[get_column_letter(index)].width = width
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def norm(value):
    return unicodedata.normalize("NFKC", value).strip().casefold()


def scalar(value, label, limit, required=False, single=False):
    if isinstance(value, bool) or (value is not None and not isinstance(value, (str, int, float))):
        raise ValueError(f"{label}单元格类型无效")
    text = str(value).strip() if value is not None else ""
    if (required and not text) or len(text) > limit:
        raise ValueError(f"{label}必填且最多 {limit} 字" if required else f"{label}最多 {limit} 字")
    if single and re.search(r"[,，、;；\r\n|]", text):
        raise ValueError(f"{label}每格只能填写单一项")
    return text


def positive(value, label):
    try:
        number = Decimal(str(value))
        if not number.is_finite() or number <= 0 or number >= Decimal("10000000000") or number.as_tuple().exponent < -8:
            raise ValueError()
        return format(number.normalize(), "f")
    except (ValueError, InvalidOperation):
        raise ValueError(f"{label}须为小于 100 亿的正数，最多 8 位小数") from None


def dimensions(values, required):
    present = [v is not None and str(v).strip() != "" for v in values]
    if not any(present) and not required:
        return ""
    if not all(present):
        raise ValueError("长、宽、高必须全部填写")
    return "*".join(positive(value, name) for value, name in zip(values, ["长", "宽", "高"]))


def location_bins(expression, remaining):
    """Unambiguous ranges only; check expansion size before allocating the list."""
    tokens = re.split(r"[,，、;；\n]", expression.replace("\r\n", "\n").replace("\r", "\n").strip().upper())
    bins = []
    for raw in tokens:
        token = raw.strip()
        if not token:
            raise ValueError("仓位列表含空项，请检查连续或末尾分隔符")
        if token == "待核仓位" or token.startswith("CL-UNKNOWN"):
            raise ValueError("不能导入系统保留的待核仓位或 CL-UNKNOWN 标识")
        if "-" not in token:
            if not re.fullmatch(r"[A-Z0-9_\u4e00-\u9fff]{1,64}", token):
                raise ValueError("单个仓位仅支持中英文、数字和下划线，最多 64 字；连字符仅用于范围")
            if len(bins) >= remaining:
                raise ValueError("所有数据行展开后合计最多 1000 个仓位（含已存在项）")
            bins.append(token)
            continue
        match = re.fullmatch(r"([A-Z_\u4e00-\u9fff-]*)([0-9]+)-([A-Z_\u4e00-\u9fff-]*)([0-9]+)", token)
        if not match:
            # Existing standard bins include A-00. A prefix without digits
            # followed by one numeric suffix cannot be mistaken for a range.
            if re.fullmatch(r"[A-Z_\u4e00-\u9fff]+-[0-9]+", token) and len(token) <= 64:
                if len(bins) >= remaining:
                    raise ValueError("所有数据行展开后合计最多 1000 个仓位（含已存在项）")
                bins.append(token)
                continue
            raise ValueError("范围须填写同前缀完整端点，例如 A1-A25 或 B01-B03")
        prefix, first, end_prefix, last = match.groups()
        if prefix != end_prefix or max(len(prefix + first), len(end_prefix + last)) > 64:
            raise ValueError("范围必须同前缀，仓位长度不能超过 64 字")
        padded = (len(first) > 1 and first.startswith("0")) or (len(last) > 1 and last.startswith("0"))
        if padded and len(first) != len(last):
            raise ValueError("带前导零的范围两端数字宽度必须一致，例如 B01-B03")
        start, end = int(first), int(last)
        if end < start:
            raise ValueError("仓位范围不能逆序")
        if end - start + 1 > remaining - len(bins):
            raise ValueError("所有数据行展开后合计最多 1000 个仓位（含已存在项）")
        bins.extend(prefix + (str(n).zfill(len(first)) if padded else str(n)) for n in range(start, end + 1))
    return bins


def parse(content, filename, kind):
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(422, "只支持 .xlsx 或 .xlsm 文件")
    if len(content) > MAX_BYTES:
        raise HTTPException(422, "文件不能超过 5 MB")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 100 or sum(i.file_size for i in archive.infolist()) > 25 * 1024 * 1024:
                raise ValueError("工作簿解压后过大")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, ValueError, OSError, KeyError, ParseError, InvalidFileException) as exc:
        raise HTTPException(422, "无法读取 Excel 工作簿，请使用下载的模板") from exc
    errors, entries = [], []
    expanded_locations = 0
    try:
        if "导入数据" not in book.sheetnames or len(book.sheetnames) > 2:
            raise HTTPException(422, "须保留“导入数据”工作表，仅允许数据与说明两张表")
        for sheet in book:
            if sheet.title not in {"导入数据", "说明与示例"}:
                raise HTTPException(422, f"不支持工作表：{sheet.title}，请使用模板")
        sheet = book["导入数据"]
        if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > len(HEADERS[kind]):
            raise HTTPException(422, "数据最多 1000 行且不能有模板外的列")
        # Do not trust cached dimensions: third-party exports can under-report them.
        sheet.reset_dimensions()
        rows = sheet.iter_rows()
        if [c.value for c in next(rows, [])] != HEADERS[kind]:
            raise HTTPException(422, "模板列名或顺序不正确，请重新下载对应模板")
        for index, cells in enumerate(rows, 2):
            if index > MAX_ROWS + 1 or len(cells) > len(HEADERS[kind]):
                raise HTTPException(422, "数据最多 1000 行且不能有模板外的列")
            values = [c.value for c in cells]
            values.extend([None] * (len(HEADERS[kind]) - len(values)))
            if all(v is None or str(v).strip() == "" for v in values):
                continue
            try:
                if any(c.data_type in {"f", "e"} for c in cells):
                    raise ValueError("不接受公式（含缓存结果）或 Excel 错误，请粘贴为值")
                if kind == "locations":
                    if len(cells) < 2 or any(c.data_type not in {"s", "inlineStr"} or not isinstance(c.value, str) for c in cells[:2]):
                        raise ValueError("仓库和仓位或范围必须是真实文本；请先设为文本后重新输入，或以单引号开头输入")
                    warehouse = scalar(values[0], "仓库", 64, True, True).upper()
                    if warehouse == "待核仓位" or warehouse.startswith("CL-UNKNOWN"):
                        raise ValueError("不能导入系统保留的待核仓位或 CL-UNKNOWN 标识")
                    if len(warehouse) > 64:
                        raise ValueError("仓库转大写后不能超过 64 字")
                    expression = scalar(values[1], "仓位或范围", 16000, True)
                    bins = location_bins(expression, MAX_ROWS - expanded_locations)
                    expanded_locations += len(bins)
                    entry = {"warehouse": warehouse, "bins": bins}
                elif kind == "paper-options":
                    entry = {"paper_types": scalar(values[0], "纸品类型", 64, single=True),
                             "paper_qualities": scalar(values[1], "纸质", 128, single=True),
                             "specifications": dimensions(values[2:5], False)}
                else:
                    if not cells or cells[0].data_type not in {"s", "inlineStr"} or not isinstance(values[0], str):
                        raise ValueError("货号必须是真实文本，不能使用数字单元格或补零显示格式；请先设为文本后重新输入，或以单引号开头输入")
                    entry = {"code": scalar(values[0], "货号", 128, True),
                             "product_name": scalar(values[1], "产品名称", 255, True),
                             "group": scalar(values[2], "配置组", 128) or "默认",
                             "line": {"packaging_type": scalar(values[3], "纸品类型", 64, True, True),
                                      "paper_quality": scalar(values[4], "纸质", 128, True, True),
                                      "specification": dimensions(values[5:8], True),
                                      "dimension_unit": scalar(values[8], "尺寸单位", 16, True),
                                      "unit": scalar(values[9], "计量单位", 32, True),
                                      "usage_quantity": positive(values[10], "每箱装产品数")},
                             "preferred": scalar(values[11], "推荐", 2) or "否",
                             "status": scalar(values[12], "状态", 2) or "启用",
                             "note": scalar(values[13], "备注", 1000)}
                    if entry["preferred"] not in {"是", "否"} or entry["status"] not in {"启用", "停用"}:
                        raise ValueError("推荐须填是/否，状态须填启用/停用")
                entries.append((index, entry))
            except ValueError as exc:
                errors.append(f"第 {index} 行：{exc}")
    except (BadZipFile, ValueError, OSError, KeyError, ParseError, InvalidFileException) as exc:
        raise HTTPException(422, "无法读取 Excel 数据页，请使用下载的模板") from exc
    finally:
        book.close()
    if not entries and not errors:
        errors.append("导入数据页没有有效数据；请将数据填入数据页，说明页不参与导入")
    return entries, errors


def snapshot(db, factory, kind):
    state = [[row.identity, master.record_out(row)] for row in db.scalars(
        select(Record).where(Record.factory_id == factory).order_by(Record.id))]
    if kind == "locations":
        state.append([[row.id, row.warehouse, row.bin_code, row.status, row.revision] for row in db.scalars(
            select(CartonLocation).where(CartonLocation.factory_id == factory).order_by(CartonLocation.id))])
    return master.digest(state)


def plan_locations(db, factory, entries, errors):
    existing = {(row.warehouse, row.bin_code): row for row in db.scalars(
        select(CartonLocation).where(CartonLocation.factory_id == factory))}
    seen, writes, details, skipped = {}, [], [], 0
    for index, entry in entries:
        for bin_code in entry["bins"]:
            key = (entry["warehouse"], bin_code)
            label = " / ".join(key)
            if key in seen:
                errors.append(f"第 {index} 行：{label} 与第 {seen[key]} 行重复")
                continue
            seen[key] = index
            row = existing.get(key)
            if row and row.id.startswith("CL-UNKNOWN"):
                errors.append(f"第 {index} 行：{label} 是系统保留仓位，不能导入")
            elif row:
                skipped += 1
                details.append(f"{label}：已存在（{'停用' if row.status == 'INACTIVE' else '启用'}），跳过")
            else:
                writes.append(key)
                details.append(f"{label}：新增启用仓位")
    return {"added": len(writes), "skipped": skipped, "errors": errors, "details": details}, writes


def plan(db, factory, kind, entries, errors):
    if kind == "locations":
        return plan_locations(db, factory, entries, errors)
    rows = list(db.scalars(select(Record).where(Record.factory_id == factory)))
    additions, skipped, writes, details = 0, 0, [], []
    if kind == "paper-options":
        row = next((r for r in rows if r.kind == "RULE" and not r.customer_code), None)
        data = json.loads(row.data_json) if row else {}
        for index, entry in entries:
            for key, value in entry.items():
                if not value:
                    continue
                hidden_key = "hidden_" + key
                visible = {norm(v) for v in data.get(key, [])}
                hidden = {norm(v) for v in data.get(hidden_key, [])}
                if norm(value) in visible and norm(value) not in hidden:
                    skipped += 1
                    continue
                if norm(value) not in visible:
                    data[key] = [*data.get(key, []), value]
                data[hidden_key] = [v for v in data.get(hidden_key, []) if norm(v) != norm(value)]
                additions += 1
                details.append(f"第 {index} 行：{value}")
        if additions:
            try:
                request = MasterSave(factory_id=factory, kind="RULE", code=row.code if row else "", data=data,
                                     status=row.status if row else "ACTIVE", preferred=bool(row.preferred) if row else False,
                                     expected_revision=row.revision if row else 0, reason="模板导入纸品选项")
                writes.append((row.id if row else "", request))
            except ValidationError:
                errors.append("纸品选项合计超出限制（类型 200 项、纸质 500 项、规格 1000 项）或已有规则无效")
    else:
        groups = {}
        for index, entry in entries:
            key = (norm(entry["code"]), norm(entry["group"]))
            group = groups.setdefault(key, {**entry, "lines": [], "index": index})
            if any(group[field] != entry[field] for field in ("code", "product_name", "group", "preferred", "status", "note")):
                errors.append(f"第 {index} 行：同一货号/配置组的产品、推荐、状态或备注冲突")
                continue
            if entry["line"] in group["lines"]:
                errors.append(f"第 {index} 行：同组纸品行重复")
            group["lines"].append(entry["line"])
        seen, preferred = set(), set()
        existing = {master.digest([r.code, {"product_name": json.loads(r.data_json).get("product_name", ""),
                    "lines": master.canonical_lines(json.loads(r.data_json).get("lines", []))}]) for r in rows if r.kind == "CONFIG"}
        identities = {r.identity for r in rows if r.kind == "CONFIG"}
        for group in groups.values():
            semantic = {"product_name": group["product_name"], "lines": master.canonical_lines(group["lines"])}
            signature = master.digest([group["code"], semantic])
            if signature in seen:
                errors.append(f"第 {group['index']} 行：不同配置组包含相同配置，请合并后导入")
            seen.add(signature)
            if group["preferred"] == "是" and group["status"] == "启用":
                if norm(group["code"]) in preferred:
                    errors.append(f"货号 {group['code']} 只能有一个启用的推荐配置")
                preferred.add(norm(group["code"]))
            if signature in existing:
                skipped += 1
                details.append(f"{group['code']} / {group['group']}：相同配置已存在，跳过")
                continue
            if signature in identities:
                errors.append(f"货号 {group['code']} 与已维护资料的原始配置冲突，请在现有资料中核对")
                continue
            try:
                request = MasterSave(factory_id=factory, kind="CONFIG", code=group["code"],
                                     data={**semantic, "note": group["note"]}, preferred=group["preferred"] == "是",
                                     status="ACTIVE" if group["status"] == "启用" else "INACTIVE", reason="模板导入货号包装")
                writes.append(("", request))
                additions += 1
                details.append(f"{group['code']} / {group['group']}：新增 {len(group['lines'])} 行纸品")
            except ValidationError:
                errors.append(f"第 {group['index']} 行：配置超出限制（每组最多 50 行纸品，请核对装箱数精度）")
    return {"added": additions, "skipped": skipped, "errors": errors, "details": details}, writes


def run(db, user, factory, kind, content, filename, expected=None, *, parsed=None):
    from app.services.carton_procurement import _lock_receipt_factory, _audit
    master.require_manage(db, user, factory)
    if expected is not None:
        _lock_receipt_factory(db, factory)
    entries, errors = parsed if parsed is not None else parse(content, filename, kind)
    fingerprint = hashlib.sha256(content).hexdigest()
    revision = snapshot(db, factory, kind)
    token = master.digest(["master-import-v1", factory, kind, fingerprint, revision])
    if expected is not None and expected != token:
        raise HTTPException(409, "文件、厂区、导入类型或基础资料已变化，请重新预览")
    result, writes = plan(db, factory, kind, entries, errors)
    result.update(factory_id=factory, kind=kind, fingerprint=fingerprint, master_revision=revision, preview_token=token)
    if expected is not None:
        if result["errors"]:
            raise HTTPException(422, "；".join(result["errors"]))
        try:
            if kind == "locations":
                from app.services import carton_positions
                for warehouse, bin_code in writes:
                    row = carton_positions.create_location(db, factory, warehouse, bin_code)
                    _audit(db, user, factory, "INVENTORY_LOCATION_CREATED", "carton_location", row.id,
                           {**carton_positions.location_out(row), "reason": "模板导入仓库仓位"})
            else:
                for identifier, request in writes:
                    master.save_record(db, user, request, identifier, commit=False)
            db.commit()
        except Exception:
            db.rollback()
            raise
    return result
