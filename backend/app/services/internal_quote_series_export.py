"""Export one selected version per product as a single series workbook."""
from copy import deepcopy
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import quote_sheetname

from app.models.internal_quote import InternalQuoteExportFile
from app.services.internal_quote import _get_quote, _check_revision, _add_audit, ensure_quote_read
from app.services.internal_quote_alternatives import issue_alternative
from app.services.internal_quote_artifacts import (
    _ensure_export_permission, create_controlled_export, safe_file_name, digest,
)
from app.services.internal_quote_calculator import canonical_json
from app.services.internal_quote_document import document_products, product_root
from app.services.internal_quote_excel import _copy_attachment_worksheet, _unique_sheet_title
from app.services.transaction_lock import lock_transaction


def _formula(value, names):
    if not isinstance(value, str) or not value.startswith("="):
        return value
    tokens = Tokenizer(value).items
    for token in tokens:
        if token.subtype != "RANGE" or "!" not in token.value:
            continue
        prefix, address = token.value.rsplit("!", 1)
        sheet = prefix[1:-1].replace("''", "'") if prefix.startswith("'") and prefix.endswith("'") else prefix
        if sheet in names:
            token.value = f"{quote_sheetname(names[sheet])}!{address}"
    return "=" + "".join(token.value for token in tokens)


def combine_series_workbooks(products):
    """Retain frozen formula dependencies privately; expose one sheet per product."""
    target = Workbook()
    target.remove(target.active)
    used = set()
    visible = []
    sources = []
    try:
        for index, (name, content) in enumerate(products, 1):
            source = load_workbook(BytesIO(content), data_only=False, keep_links=False)
            sources.append(source)
            names = {sheet.title: _unique_sheet_title(
                name if offset == 0 else f"{index}-{sheet.title}", used,
            ) for offset, sheet in enumerate(source.worksheets)}
            for offset, sheet in enumerate(source.worksheets):
                copied = target.create_sheet(names[sheet.title])
                _copy_attachment_worksheet(sheet, copied, {})
                for row in copied:
                    for cell in row:
                        if cell.data_type == "f":
                            cell.value = _formula(cell.value, names)
                copied.print_area = _formula("=" + str(sheet.print_area), names)[1:] if sheet.print_area else None
                copied.print_title_rows = sheet.print_title_rows
                copied.print_title_cols = sheet.print_title_cols
                copied.row_breaks = deepcopy(sheet.row_breaks)
                copied.col_breaks = deepcopy(sheet.col_breaks)
                copied.sheet_state = "visible" if offset == 0 else "veryHidden"
                if offset == 0:
                    visible.append(copied)
        target._sheets = visible + [sheet for sheet in target.worksheets if sheet not in visible]
        target.active = 0
        target.calculation.fullCalcOnLoad = True
        target.calculation.forceFullCalc = True
        target.calculation.calcMode = "auto"
        buffer = BytesIO()
        target.save(buffer)
        return buffer.getvalue()
    finally:
        target.close()
        for source in sources:
            source.close()


def export_series(db, quote_id, payload, user, request=None):
    """Freeze all unissued members atomically; never regenerate issued bytes."""
    try:
        ids = [row.quote_id for row in payload.products]
        if len(ids) != len(set(ids)):
            raise HTTPException(400, "每款产品只能选择一个报价版本")
        source = _get_quote(db, quote_id)
        quotes = [_get_quote(db, identity) for identity in ids]
        for quote in [source, *quotes]:
            ensure_quote_read(db, user, quote.factory_id)
            _ensure_export_permission(db, quote, user)
        roots = [product_root(db, quote) for quote in [source, *quotes]]
        for identity in sorted({f"{q.factory_id}:{q.batch_id or q.id}" for q in [source, *quotes, *roots]}):
            lock_transaction(db, "internal-quote", identity)
        for identity in sorted({root.id for root in roots}):
            lock_transaction(db, "internal-quote-family", identity)
        db.expire_all()
        source = _get_quote(db, quote_id)
        quotes = [_get_quote(db, identity) for identity in ids]
        members = document_products(db, source)
        expected = {row.id for row in members}
        selected = [product_root(db, quote).id for quote in quotes]
        if len(selected) != len(expected) or set(selected) != expected or len(set(selected)) != len(selected):
            raise HTTPException(400, "请选择本系列的每款产品，且每款只选择一个版本")
        revisions = {row.quote_id: row.revision for row in payload.products}
        by_root = dict(zip(selected, quotes))
        ordered = [by_root[row.id] for row in members]
        for quote in ordered:
            if quote.module_version != "v4" and not (
                quote.status in {"fully_approved", "exported"}
                and quote.final_release_status == "approved" and quote.final_release_revision > 0
            ):
                raise HTTPException(409, f"{quote.product_name}：历史报价须通过最终审核后才能系列输出")
        records = []
        for quote in ordered:
            if quote.factory_id != source.factory_id or quote.customer != source.customer:
                raise HTTPException(400, "系列报价必须属于同一客户和厂区")
            if quote.module_version != "v4" or quote.final_release_status != "issued":
                _check_revision(quote.header_revision, revisions[quote.id], quote.product_name)
            try:
                if quote.module_version == "v4":
                    out = issue_alternative(db, quote.id, next(row for row in payload.products if row.quote_id == quote.id), user, request, commit=False)
                else:
                    out = create_controlled_export(db, quote.id, user, request, commit=False)
            except HTTPException as error:
                raise HTTPException(error.status_code, f"{quote.product_name}：{error.detail}") from error
            records.append(db.get(InternalQuoteExportFile, out.id))
        content = combine_series_workbooks([(quote.product_name, record.content) for quote, record in zip(ordered, records)])
        sha256 = digest(content)
        _add_audit(db, source, user, "export_series", detail=canonical_json({
            "sha256": sha256,
            "products": [{"quote_id": quote.id, "export_id": record.id, "sha256": record.sha256}
                         for quote, record in zip(ordered, records)],
        }), request=request)
        db.commit()
        return content, safe_file_name(f"{members[0].batch_quote_no or members[0].quote_no}-系列报价.xlsx"), sha256
    except Exception:
        db.rollback()
        raise
