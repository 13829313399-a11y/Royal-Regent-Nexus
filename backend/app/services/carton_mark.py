from __future__ import annotations

import re
import os
import shutil
import string
import subprocess
import unicodedata
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from dataclasses import dataclass
from difflib import SequenceMatcher
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from app.schemas.carton_mark import (
    CartonMarkAutoCheckResponse,
    CartonMarkAutoCheckSummary,
    CartonMarkBatchCheckItem,
    CartonMarkBatchCheckResponse,
    CartonMarkComparisonItem,
    CartonMarkDocumentCheckResponse,
    CartonMarkDocumentCheckSummary,
    CartonMarkDocumentComparisonItem,
    CartonMarkDocumentTextItem,
    CartonMarkExtractedField,
    CartonMarkExtractionStatus,
)
from app.services.carton_mark_document_scope import prepare_carton_mark_document_scope


@dataclass(frozen=True)
class FieldDefinition:
    key: str
    label: str
    aliases: tuple[str, ...]


@dataclass(frozen=True)
class PdfMarkRegion:
    kind: str
    text: str
    box: tuple[int, int, int, int]
    # Zero-based PDF page number.  Keep this separate from ``box`` so page
    # selection never has to infer a page from synthetic coordinate offsets.
    page_index: int = 0


@dataclass(frozen=True)
class PdfTextSpan:
    text: str
    x: float
    y: float


@dataclass
class _PdfVectorDocumentLine:
    text: str
    x: float
    y: float
    order: int
    field_key: str
    label_only: bool
    starts_field_label: bool


@dataclass(frozen=True)
class OcrWord:
    text: str
    left: int
    top: int
    right: int
    bottom: int
    confidence: float


@dataclass(frozen=True)
class DocumentExtraction:
    items: list[CartonMarkDocumentTextItem]
    status: CartonMarkExtractionStatus
    reviews: list[tuple[str, str]]


@dataclass(frozen=True)
class DocumentCharacter:
    value: str
    item_index: int
    original_index: int
    break_before: bool


FIELD_DEFINITIONS = [
    FieldDefinition("customer_name", "客名", ("CUSTOMER", "CLIENT", "CLIENTE", "NOMBRE", "CUSTOMER NAME")),
    FieldDefinition("po", "PO", (
        "PO", "P.O.", "P/O", "PO NO", "PO NO.", "P.O. NO", "P.O.NO", "PO NUMBER", "ORDER NO",
        "ORDER NO.", "ORDER NUMBER", "NUMERO DE PEDIDO", "NÚMERO DE PEDIDO", "PEDIDO",
    )),
    FieldDefinition("pi_no", "PI No", (
        "PI NO", "PI NO.", "P.I. NO", "P.I. NO.", "PINO", "PINO.",
        "PROFORMA INVOICE", "PROFORMA INVOICE NO", "INVOICE NO", "INVOICE NO.",
    )),
    FieldDefinition("item", "ITEM", (
        "ITEM", "ITEM NO", "ITEM NO.", "ITEM NUMBER", "ITEM#", "STYLE", "STYLE NO", "STYLE NO.",
        "MODEL", "MODELO", "ART NO", "ART NO.", "ARTICLE", "ARTICLE NO", "ARTICLE NO.",
        "ARTICLE NUMBER", "REF", "REFERENCE",
    )),
    FieldDefinition("description", "描述", ("DESCRIPTION", "DESCRIPCION", "DESCRIPCIÓN")),
    FieldDefinition("sku", "SKU", ("SKU", "S.K.U.", "SKU NO", "SKU NO.")),
    FieldDefinition("color", "颜色", ("COLOR", "COLOUR", "COL")),
    FieldDefinition("quantity", "数量", (
        "QTY", "QUANTITY", "CANTIDAD", "PCS", "PZAS", "CTN QTY", "QTY/CTN", "QTY PER CTN",
        "QTY PER CARTON", "PCS/CTN", "PIEZAS POR BULTO",
        "NUMBER OF PIECES", "NUMBEROFPIECES",
    )),
    FieldDefinition("carton_no", "箱号", (
        "CARTON NO", "CARTON NO.", "CARTON NUMBER", "CTN NO", "CTN NO.", "CTN#", "C/NO",
        "BULTO", "BULTOS", "BULTO NO",
        "BULTO NO.", "NO DE BULTO", "NO. DE BULTO", "NRO BULTO", "NRO. BULTO",
    )),
    FieldDefinition("box_no", "箱序", ("CAJA NUMERO", "CAJA NÚMERO", "BOX NO", "BOX NUMBER")),
    FieldDefinition("size", "尺码", ("SIZE", "TALLA")),
    FieldDefinition("gw", "G.W", (
        "G.W", "G.W.", "GW", "GROSS WEIGHT", "CARTON GROSS WEIGHT", "PESO BRUTO",
        "GRS.WT", "GRS WT", "W.G.R", "W. G.R.", "WGR",
    )),
    FieldDefinition("nw", "N.W", (
        "N.W", "N.W.", "NW", "NET WEIGHT", "CARTON NET WEIGHT", "PESO NETO",
        "W.NET", "W. NET", "W NET",
    )),
    FieldDefinition("measurement", "尺寸", (
        "MEAS", "MEAS.", "MEASUREMENT", "DIMENSION", "CARTON SIZE", "CARTON MEAS",
        "MEDIDA",
    )),
    FieldDefinition("section", "分区", ("SECCION / UNECO", "SECCIÓN / UNECO", "SECCION/UNECO", "SECCIÓN/UNECO", "SECCION", "SECCIÓN")),
    FieldDefinition("barcode", "条码", (
        "BARCODE", "BAR CODE", "BARCODE NO", "BAR CODE NO", "CODE", "EAN", "UPC",
        "CODIGO DE BARRAS", "CÓDIGO DE BARRAS",
    )),
    FieldDefinition("wall_type", "纸板层型", ("WALL TYPE", "BOARD TYPE", "WALL CONSTRUCTION")),
    FieldDefinition("bursting_test", "耐破强度", (
        "BURSTING TEST", "BURST TEST", "BUASTING TEST",
    )),
    FieldDefinition("min_combined_weight", "面纸最低克重", (
        "MIN COMB WT FACINGS", "MIN. COMB. WT. FACINGS", "MIN COMBINED WEIGHT FACINGS",
        "MIN COMB WT", "MINIMUM COMBINED WEIGHT",
    )),
    FieldDefinition("size_limit", "尺寸上限", ("SIZE LIMIT", "SIZE LT")),
    FieldDefinition("gross_weight_limit", "毛重上限", (
        "GROSS WT LT", "GROSS WT LIMIT", "GROSS WEIGHT LIMIT", "MAX GROSS WEIGHT",
    )),
]

FIELD_BY_KEY = {field.key: field for field in FIELD_DEFINITIONS}
FIELD_ORDER = [field.key for field in FIELD_DEFINITIONS]
IDENTIFIER_FIELD_KEYS = {"po", "pi_no", "item", "sku", "carton_no", "barcode"}
OCR_IDENTIFIER_TRANSLATION = str.maketrans({
    "O": "0",
    "Q": "0",
    "D": "0",
    "I": "1",
    "L": "1",
    "|": "1",
    "S": "5",
    "Z": "2",
    "B": "8",
    "G": "6",
})
TEXT_CHARS = set(string.printable) | set("，。：；、（）【】《》±×")
PROJECT_ROOT = Path(__file__).resolve().parents[3]
TESSERACT_CANDIDATE_PATHS = (
    PROJECT_ROOT / "tools" / "Tesseract-OCR" / "tesseract.exe",
    Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
    Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
)
DOCUMENT_EXCEL_SUFFIXES = {".xls", ".xlsx", ".xlsm"}
DOCUMENT_PDF_SUFFIXES = {".pdf"}
MAX_DOCUMENT_FILE_BYTES = 20 * 1024 * 1024
MAX_DOCUMENT_ZIP_ENTRIES = 20_000
MAX_DOCUMENT_ZIP_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_DOCUMENT_WORKSHEETS = 50
MAX_DOCUMENT_SHEET_ROWS = 10_000
MAX_DOCUMENT_SHEET_COLUMNS = 512
MAX_DOCUMENT_SCANNED_CELLS = 250_000
MAX_DOCUMENT_TEXT_ITEMS = 20_000
MAX_DOCUMENT_ITEM_CHARS = 4_000
MAX_DOCUMENT_TOTAL_TEXT_CHARS = 1_000_000
MAX_DOCUMENT_PDF_PAGES = 100
MAX_DOCUMENT_EXACT_SEARCH_STEPS = 250_000
MAX_DOCUMENT_FUZZY_CHECKS = 100_000
XLS_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")
DIRECT_OOXML_REFERENCE_RE = re.compile(
    r"^\s*=\s*(?:(?:'(?P<quoted>(?:[^']|'')+)'|(?P<plain>[A-Z0-9_.\u4e00-\u9fff]+))!)?"
    r"\$?(?P<column>[A-Z]{1,4})\$?(?P<row>\d+)\s*$",
    re.IGNORECASE,
)
CANONICAL_DOCUMENT_PUNCTUATION = str.maketrans({
    "\u2010": "-",
    "\u2011": "-",
    "\u2012": "-",
    "\u2013": "-",
    "\u2014": "-",
    "\u2212": "-",
    "\u2018": "'",
    "\u2019": "'",
    "\u201c": '"',
    "\u201d": '"',
})
DOCUMENT_LAYOUT_SEPARATORS = {
    ":",
    "|",
    "·",
    "•",
    "‧",
    "∙",
    "⋅",
    "◦",
    "●",
    "○",
    "▪",
    "▫",
    "■",
    "□",
    "◆",
    "◇",
    "►",
    "▸",
    "▶",
}


class CartonMarkDocumentError(ValueError):
    pass


class CartonMarkDocumentConfigurationError(RuntimeError):
    pass


def validate_carton_mark_document_file(
    file_name: str,
    content: bytes,
    *,
    kind: str,
) -> str:
    safe_name = Path(file_name or "").name
    suffix = Path(safe_name).suffix.lower()
    allowed = DOCUMENT_EXCEL_SUFFIXES if kind == "excel" else DOCUMENT_PDF_SUFFIXES
    label = "客户 Excel" if kind == "excel" else "打印 PDF"
    if not safe_name or suffix not in allowed:
        allowed_text = "/".join(sorted(allowed))
        raise CartonMarkDocumentError(f"{label}文件类型无效，仅支持 {allowed_text}")
    if not content:
        raise CartonMarkDocumentError(f"{label}不能为空")
    if len(content) > MAX_DOCUMENT_FILE_BYTES:
        raise CartonMarkDocumentError(f"{label}不能超过 20 MB")

    if kind == "pdf":
        if not content[:1024].lstrip().startswith(b"%PDF-"):
            raise CartonMarkDocumentError("打印 PDF 内容与 .pdf 扩展名不符")
        return safe_name

    if suffix == ".xls":
        if not content.startswith(XLS_MAGIC):
            raise CartonMarkDocumentError("客户 Excel 内容不是有效的旧版 .xls 工作簿")
        return safe_name

    if not content.startswith(b"PK"):
        raise CartonMarkDocumentError("客户 Excel 内容与 OOXML 扩展名不符")
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            names = {info.filename.replace("\\", "/") for info in infos}
            if len(infos) > MAX_DOCUMENT_ZIP_ENTRIES:
                raise CartonMarkDocumentError("客户 Excel 压缩包条目过多")
            if any(info.flag_bits & 0x1 for info in infos):
                raise CartonMarkDocumentError("客户 Excel 已加密，无法核对显示文字")
            if sum(info.file_size for info in infos) > MAX_DOCUMENT_ZIP_UNCOMPRESSED_BYTES:
                raise CartonMarkDocumentError("客户 Excel 解压后内容过大")
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise CartonMarkDocumentError("客户 Excel 缺少工作簿核心文件")
            for name in names:
                path = Path(name)
                if path.is_absolute() or ".." in path.parts:
                    raise CartonMarkDocumentError("客户 Excel 包含不安全的内部路径")
    except CartonMarkDocumentError:
        raise
    except (BadZipFile, OSError, ValueError) as exc:
        raise CartonMarkDocumentError("客户 Excel 不是有效的 OOXML 工作簿") from exc
    return safe_name


def build_carton_mark_document_check(
    *,
    excel_file_name: str,
    excel_bytes: bytes,
    pdf_file_name: str,
    pdf_bytes: bytes,
) -> CartonMarkDocumentCheckResponse:
    excel_name = validate_carton_mark_document_file(
        excel_file_name,
        excel_bytes,
        kind="excel",
    )
    pdf_name = validate_carton_mark_document_file(
        pdf_file_name,
        pdf_bytes,
        kind="pdf",
    )
    excel = extract_excel_document_items(excel_name, excel_bytes)
    if not excel.items and not excel.reviews:
        raise CartonMarkDocumentError("客户 Excel 没有可核对的可见非空单元格文字")
    pdf = extract_pdf_document_items(pdf_bytes)
    if sum(len(item.text) for item in excel.items) > MAX_DOCUMENT_TOTAL_TEXT_CHARS:
        raise CartonMarkDocumentError("客户 Excel 的可核对显示文字总量过大")
    if sum(len(item.text) for item in pdf.items) > MAX_DOCUMENT_TOTAL_TEXT_CHARS:
        raise CartonMarkDocumentError("打印 PDF 的可核对文字总量过大")

    scope = prepare_carton_mark_document_scope(excel.items, pdf.items)
    scope_message = (
        f"正文选区置信度 {scope.confidence:.0%}。{scope.diagnostics.describe()}"
    )
    scope_status = CartonMarkExtractionStatus(
        source="document_scope",
        ok=not scope.requires_review,
        engine="evidence-based-document-scope",
        message=scope.review_note or scope_message,
        raw_text=clip_text("\n".join(item.text for item in scope.excel_items)),
        match_confidence=scope.confidence,
        requires_review=scope.requires_review,
        review_reason=scope.review_note,
    )

    comparisons: list[CartonMarkDocumentComparisonItem]
    if scope.pdf_items:
        comparisons = compare_document_text_items(scope.excel_items, scope.pdf_items)
    else:
        comparisons = []

    for location, note in (*excel.reviews, *pdf.reviews):
        comparisons.append(CartonMarkDocumentComparisonItem(
            status="review",
            expected_location=location,
            note=note,
        ))

    if scope.requires_review:
        comparisons.append(CartonMarkDocumentComparisonItem(
            status="review",
            expected_location="Excel 箱唛正文选区",
            note=scope.review_note or "无法可靠确定 Excel 箱唛正文选区，请人工复核。",
        ))

    if not scope.pdf_items and not pdf.reviews:
        comparisons.append(CartonMarkDocumentComparisonItem(
            status="review",
            note=pdf.status.message or "打印 PDF 未能提取可核对文字，请人工复核。",
        ))

    summary = summarize_document_comparisons(
        comparisons,
        force_review=scope.requires_review or _pdf_document_requires_review(pdf),
    )
    return CartonMarkDocumentCheckResponse(
        excel_file_name=excel_name,
        pdf_file_name=pdf_name,
        summary=summary,
        excel_items=scope.excel_items,
        pdf_items=scope.pdf_items,
        comparisons=comparisons,
        extraction=[excel.status, pdf.status, scope_status],
    )


def extract_excel_document_items(file_name: str, content: bytes) -> DocumentExtraction:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".xls":
        return _extract_xls_document_items(content)
    return _extract_ooxml_document_items(content)


def _extract_ooxml_document_items(content: bytes) -> DocumentExtraction:
    try:
        from openpyxl import load_workbook  # type: ignore
        from openpyxl.utils import get_column_letter  # type: ignore
    except Exception as exc:
        raise CartonMarkDocumentConfigurationError(
            "服务器未配置 OOXML Excel 读取组件"
        ) from exc

    try:
        value_book = load_workbook(
            BytesIO(content),
            read_only=False,
            data_only=True,
            keep_links=False,
        )
        formula_book = load_workbook(
            BytesIO(content),
            read_only=False,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise CartonMarkDocumentError(f"客户 Excel 无法读取：{exc}") from exc

    items: list[CartonMarkDocumentTextItem] = []
    reviews: list[tuple[str, str]] = []
    scanned_cells = 0
    try:
        visible_sheets = [sheet for sheet in value_book.worksheets if sheet.sheet_state == "visible"]
        if len(visible_sheets) > MAX_DOCUMENT_WORKSHEETS:
            raise CartonMarkDocumentError("客户 Excel 的可见工作表过多")
        sheet_details = []
        for sheet in visible_sheets:
            if sheet.max_row > MAX_DOCUMENT_SHEET_ROWS or sheet.max_column > MAX_DOCUMENT_SHEET_COLUMNS:
                raise CartonMarkDocumentError(
                    f"工作表 {sheet.title} 的有效范围过大，无法安全核对"
                )
            sheet_cells = sheet.max_row * sheet.max_column
            scanned_cells += sheet_cells
            if scanned_cells > MAX_DOCUMENT_SCANNED_CELLS:
                raise CartonMarkDocumentError("客户 Excel 的单元格范围过大，无法安全核对")
            formula_sheet = formula_book[sheet.title]
            hidden_columns = {
                index
                for index in range(1, sheet.max_column + 1)
                if sheet.column_dimensions[get_column_letter(index)].hidden
            }
            sheet_details.append((sheet, formula_sheet, hidden_columns))

        referenced_formula_cells: set[tuple[str, str]] = set()
        for sheet, formula_sheet, hidden_columns in sheet_details:
            for row_index in range(1, sheet.max_row + 1):
                if sheet.row_dimensions[row_index].hidden:
                    continue
                for column_index in range(1, sheet.max_column + 1):
                    if column_index in hidden_columns:
                        continue
                    formula = formula_sheet.cell(row_index, column_index).value
                    if not isinstance(formula, str) or not formula.startswith("="):
                        continue
                    identity = _direct_ooxml_reference_identity(sheet.title, formula)
                    if identity is not None:
                        referenced_formula_cells.add(identity)

        for sheet, formula_sheet, hidden_columns in sheet_details:
            for row_index in range(1, sheet.max_row + 1):
                if sheet.row_dimensions[row_index].hidden:
                    continue
                for column_index in range(1, sheet.max_column + 1):
                    if column_index in hidden_columns:
                        continue
                    cell = sheet.cell(row_index, column_index)
                    formula_cell = formula_sheet.cell(row_index, column_index)
                    location = f"{sheet.title}!{cell.coordinate}"
                    value = cell.value
                    formula = formula_cell.value
                    if value is None and isinstance(formula, str) and formula.startswith("="):
                        value = _resolve_direct_ooxml_reference_value(
                            value_book,
                            formula_book,
                            sheet.title,
                            formula,
                        )
                        if value is None:
                            reviews.append((
                                location,
                                "该公式单元格没有已保存的显示结果，服务器不会执行公式，请在 Excel 中重新计算并保存后再核对。",
                            ))
                            continue
                    if (
                        isinstance(formula, str)
                        and _direct_ooxml_reference_identity(sheet.title, formula) is not None
                        and (sheet.title, cell.coordinate) in referenced_formula_cells
                    ):
                        # Intermediate mirror cells are controls feeding another
                        # visible formula.  The downstream printable mirror keeps
                        # the same source wording; counting both creates a false
                        # third occurrence beside the front/side PDF panels.
                        continue
                    text = display_excel_cell_value(value, cell.number_format)
                    _append_document_item(items, text, location)
    finally:
        value_book.close()
        formula_book.close()

    status_message = ""
    if reviews:
        status_message = f"有 {len(reviews)} 个公式单元格没有已保存的显示结果。"
    return DocumentExtraction(
        items=items,
        status=CartonMarkExtractionStatus(
            source="customer_excel",
            ok=bool(items),
            engine="openpyxl",
            message=status_message,
            raw_text=clip_text("\n".join(item.text for item in items)),
        ),
        reviews=reviews,
    )


def _resolve_direct_ooxml_reference_value(
    value_book,
    formula_book,
    current_sheet: str,
    formula: str,
    *,
    seen: frozenset[tuple[str, str]] = frozenset(),
    depth: int = 0,
):
    """Resolve only a direct cell mirror, never execute an Excel expression.

    Customer carton templates often use ``=B17`` to mirror a source-table value
    into the printable mark.  openpyxl deliberately does not calculate formulas,
    but following an exact cell reference is deterministic and avoids a needless
    manual review.  Functions, arithmetic, external links and cycles stay review.
    """

    # Direct mirrors are deterministic, but a malicious or simply malformed
    # workbook can contain an arbitrarily long acyclic chain.  Bound recursion
    # before resolving the next target so such a chain becomes the existing
    # formula-review path instead of raising RecursionError.
    if depth >= 128:
        return None

    identity = _direct_ooxml_reference_identity(current_sheet, formula)
    if identity is None:
        return None
    sheet_name, coordinate = identity
    if sheet_name not in value_book.sheetnames or sheet_name not in formula_book.sheetnames:
        return None
    if identity in seen:
        return None

    target_value = value_book[sheet_name][coordinate].value
    if target_value is not None:
        return target_value
    target_formula = formula_book[sheet_name][coordinate].value
    if isinstance(target_formula, str) and target_formula.startswith("="):
        return _resolve_direct_ooxml_reference_value(
            value_book,
            formula_book,
            sheet_name,
            target_formula,
            seen=seen | {identity},
            depth=depth + 1,
        )
    return None


def _direct_ooxml_reference_identity(
    current_sheet: str,
    formula: str,
) -> tuple[str, str] | None:
    match = DIRECT_OOXML_REFERENCE_RE.fullmatch(str(formula or ""))
    if not match:
        return None
    sheet_name = match.group("quoted") or match.group("plain") or current_sheet
    sheet_name = sheet_name.replace("''", "'")
    coordinate = f"{match.group('column').upper()}{match.group('row')}"
    return sheet_name, coordinate


def _extract_xls_document_items(content: bytes) -> DocumentExtraction:
    try:
        import xlrd  # type: ignore
    except Exception as exc:
        raise CartonMarkDocumentConfigurationError(
            "服务器未配置旧版 .xls 读取组件"
        ) from exc

    try:
        workbook = xlrd.open_workbook(file_contents=content, formatting_info=True)
    except Exception as exc:
        raise CartonMarkDocumentError(f"客户 .xls 无法读取：{exc}") from exc
    if workbook.nsheets > MAX_DOCUMENT_WORKSHEETS:
        raise CartonMarkDocumentError("客户 Excel 的工作表过多")

    items: list[CartonMarkDocumentTextItem] = []
    scanned_cells = 0
    for sheet in workbook.sheets():
        if getattr(sheet, "visibility", 0):
            continue
        if sheet.nrows > MAX_DOCUMENT_SHEET_ROWS or sheet.ncols > MAX_DOCUMENT_SHEET_COLUMNS:
            raise CartonMarkDocumentError(f"工作表 {sheet.name} 的有效范围过大，无法安全核对")
        scanned_cells += sheet.nrows * sheet.ncols
        if scanned_cells > MAX_DOCUMENT_SCANNED_CELLS:
            raise CartonMarkDocumentError("客户 Excel 的单元格范围过大，无法安全核对")
        for row_index in range(sheet.nrows):
            if getattr(sheet.rowinfo_map.get(row_index), "hidden", False):
                continue
            for column_index in range(sheet.ncols):
                if getattr(sheet.colinfo_map.get(column_index), "hidden", False):
                    continue
                cell = sheet.cell(row_index, column_index)
                value = cell.value
                if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                    continue
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate_as_datetime(value, workbook.datemode)
                    except (TypeError, ValueError, OverflowError):
                        pass
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = bool(value)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    value = xlrd.error_text_from_code.get(value, f"#{value}")

                number_format = "General"
                try:
                    xf = workbook.xf_list[cell.xf_index]
                    number_format = workbook.format_map[xf.format_key].format_str
                except (AttributeError, IndexError, KeyError):
                    pass
                location = f"{sheet.name}!{xlrd.formula.cellname(row_index, column_index)}"
                _append_document_item(
                    items,
                    display_excel_cell_value(value, number_format),
                    location,
                )

    return DocumentExtraction(
        items=items,
        status=CartonMarkExtractionStatus(
            source="customer_excel",
            ok=bool(items),
            engine="xlrd",
            message="",
            raw_text=clip_text("\n".join(item.text for item in items)),
        ),
        reviews=[],
    )


def _append_document_item(
    items: list[CartonMarkDocumentTextItem],
    text: str,
    location: str,
) -> None:
    text = str(text or "").replace("\x00", "").strip()
    if not document_canonical_value(text):
        return
    if len(text) > MAX_DOCUMENT_ITEM_CHARS:
        raise CartonMarkDocumentError(f"{location} 的显示文字过长，无法安全核对")
    if len(items) >= MAX_DOCUMENT_TEXT_ITEMS:
        raise CartonMarkDocumentError("客户 Excel 的非空文字单元格过多")
    items.append(CartonMarkDocumentTextItem(text=text, location=location))


def display_excel_cell_value(value, number_format: str = "General") -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return _format_excel_datetime(value, number_format)
    if isinstance(value, date):
        return _format_excel_datetime(datetime.combine(value, time()), number_format)
    if isinstance(value, time):
        return _format_excel_datetime(datetime.combine(date(1899, 12, 30), value), number_format)
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return _format_excel_number(value, number_format)
    return str(value).strip()


def _format_excel_number(value, number_format: str) -> str:
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return str(value).strip()
    if not number.is_finite():
        return str(value).strip()

    format_text = str(number_format or "General").split(";", 1)[0]
    simple_format = re.sub(r"\[[^\]]*\]", "", format_text)
    simple_format = re.sub(r'"([^"]*)"', r"\1", simple_format).strip()
    if "%" in simple_format:
        number *= 100
    decimal_match = re.search(r"[0#?,]+\.([0#?]+)", simple_format)
    decimal_places = len(decimal_match.group(1)) if decimal_match else 0
    integer_pattern = simple_format.split(".", 1)[0]
    zero_width = len(re.findall(r"0", integer_pattern))
    use_grouping = "," in integer_pattern

    if decimal_match:
        rendered = f"{number:,.{decimal_places}f}" if use_grouping else f"{number:.{decimal_places}f}"
    elif number == number.to_integral_value():
        rendered = str(int(number))
        if zero_width > 1 and not use_grouping:
            sign = "-" if rendered.startswith("-") else ""
            digits = rendered.lstrip("-").zfill(zero_width)
            rendered = sign + digits
        elif use_grouping:
            rendered = f"{int(number):,}"
    else:
        rendered = format(number.normalize(), "f")
    if "%" in simple_format:
        rendered += "%"
    return rendered


def _format_excel_datetime(value: datetime, number_format: str) -> str:
    fmt = str(number_format or "").lower()
    fmt = re.sub(r"\[[^\]]*\]", "", fmt)
    if not fmt or fmt == "general":
        rendered = value.isoformat(sep=" ", timespec="seconds")
        return rendered[:-9] if rendered.endswith(" 00:00:00") else rendered

    has_date = bool(re.search(r"[yd]", fmt))
    has_time = bool(re.search(r"h|s", fmt))
    if has_date:
        year = f"{value.year:04d}" if "yyyy" in fmt else f"{value.year % 100:02d}"
        month = f"{value.month:02d}" if "mm" in fmt else str(value.month)
        day = f"{value.day:02d}" if "dd" in fmt else str(value.day)
        token_positions = {
            "y": min((pos for token in ("yyyy", "yy") if (pos := fmt.find(token)) >= 0), default=10**6),
            "m": min((pos for token in ("mmmm", "mmm", "mm", "m") if (pos := fmt.find(token)) >= 0), default=10**6),
            "d": min((pos for token in ("dddd", "ddd", "dd", "d") if (pos := fmt.find(token)) >= 0), default=10**6),
        }
        order = [key for key, _ in sorted(token_positions.items(), key=lambda item: item[1])]
        values = {"y": year, "m": month, "d": day}
        if "年" in fmt or "月" in fmt or "日" in fmt:
            date_text = f"{year}年{month}月{day}日"
        else:
            separator_match = re.search(r"[ymd]+([^ymd\w]+)[ymd]+", fmt)
            separator = separator_match.group(1).replace("\\", "") if separator_match else "-"
            separator = separator or "-"
            date_text = separator.join(values[key] for key in order)
    else:
        date_text = ""

    if has_time:
        hour = value.hour
        suffix = ""
        if "am/pm" in fmt:
            suffix = " AM" if hour < 12 else " PM"
            hour = hour % 12 or 12
        hour_text = f"{hour:02d}" if "hh" in fmt else str(hour)
        minute_text = f"{value.minute:02d}"
        second_text = f":{value.second:02d}" if "s" in fmt else ""
        time_text = f"{hour_text}:{minute_text}{second_text}{suffix}"
    else:
        time_text = ""
    return " ".join(part for part in (date_text, time_text) if part)


def _clean_pdf_document_line(value: str) -> str:
    """Drop broken embedded-font output without losing ordinary business text."""

    cleaned: list[str] = []
    script_counts: dict[str, int] = {}
    supported_letters = 0
    unsupported_letters = 0
    for character in str(value or ""):
        category = unicodedata.category(character)
        if (
            character in {"\ufffd", "\ufffc", "\ufeff", "\u00ad"}
            or category in {"Cc", "Cf", "Cs", "Co", "Cn", "Zl", "Zp"}
        ):
            cleaned.append(" ")
            continue
        cleaned.append(character)
        name = unicodedata.name(character, "")
        script = next((
            candidate
            for candidate in (
                "LATIN", "CJK", "HIRAGANA", "KATAKANA", "HANGUL", "CYRILLIC",
                "ARABIC", "GREEK", "HEBREW", "DEVANAGARI", "BENGALI", "GUJARATI",
                "GURMUKHI", "TAMIL", "TELUGU", "KANNADA", "MALAYALAM", "THAI",
                "LAO", "GEORGIAN", "ARMENIAN", "TIBETAN", "MYANMAR",
            )
            if candidate in name
        ), "")
        if not character.isalpha() and not (category.startswith("M") and script):
            continue
        if script:
            script_counts[script] = script_counts.get(script, 0) + 1
            supported_letters += 1
        else:
            unsupported_letters += 1

    text = re.sub(r"\s+", " ", "".join(cleaned)).strip()
    letter_count = supported_letters + unsupported_letters
    if not document_canonical_value(text) or not any(character.isalnum() for character in text):
        return ""
    # Legitimate carton text in this workflow may mix Latin with Chinese,
    # Cyrillic or Arabic.  Broken subset fonts instead produce several unrelated
    # scripts (Tibetan, Georgian, Hebrew, Indic glyphs, etc.) in one short line.
    if unsupported_letters > max(1, int(letter_count * 0.12)):
        return ""
    significant_scripts = sum(count >= 2 for count in script_counts.values())
    if len(script_counts) >= 3 or significant_scripts > 2:
        return ""
    if (
        len(script_counts) >= 2
        and letter_count <= 4
        and "LATIN" not in script_counts
    ):
        return ""
    compact = document_canonical_value(text)
    if text.upper() in {"PAP", "RESY"}:
        return ""
    if (
        len([character for character in text if character.isalnum()]) <= 2
        and any(ord(character) > 127 and "LATIN" in unicodedata.name(character, "") for character in text)
    ):
        return ""
    return text


_PDF_VECTOR_FIELD_PREFIXES: tuple[tuple[str, str], ...] = (
    # More specific carton/table labels must precede their shorter aliases.
    ("CARTONNETWEIGHT", "nw"),
    ("CARTONGROSSWEIGHT", "gw"),
    ("CARTONSIZECM", "measurement"),
    ("CARTONMEASUREMENT", "measurement"),
    ("CARTONSIZE", "measurement"),
    ("PRODUCTDESCRIPTION", "description"),
    ("ITEMDESCRIPTION", "description"),
    ("NUMBEROFPIECES", "quantity"),
    ("ARTICLEQUANTITY", "quantity"),
    ("QTYPERCARTON", "quantity"),
    ("QTYPERCTN", "quantity"),
    ("PIEZASPORBULTO", "quantity"),
    ("CARTONNUMBER", "carton_no"),
    ("CARTONNO", "carton_no"),
    ("CTNNUMBER", "carton_no"),
    ("CTNNO", "carton_no"),
    ("CARTONCBM", ""),
    ("BARCODENUMBER", "barcode"),
    ("BARCODENO", "barcode"),
    ("BARCODE", "barcode"),
    ("BARCODE", "barcode"),
    ("ITEMNUMBER", "item"),
    ("ITEMNO", "item"),
    ("ARTICLENUMBER", "item"),
    ("ARTICLENO", "item"),
    ("DESCRIPTION", "description"),
    ("CTNQTY", "quantity"),
    ("QTYCTN", "quantity"),
    ("PCSCTN", "quantity"),
    ("QUANTITY", "quantity"),
    ("NETWEIGHT", "nw"),
    ("NETWT", "nw"),
    ("GROSSWEIGHT", "gw"),
    ("GROSSWT", "gw"),
    ("MEASUREMENT", "measurement"),
    ("MEAS", "measurement"),
    ("WNET", "nw"),
    ("WGR", "gw"),
    ("GRSWT", "gw"),
    ("PINO", "pi_no"),
    ("ORDERNUMBER", "po"),
    ("ORDERNO", "po"),
    ("PONUMBER", "po"),
    ("PONO", "po"),
    ("BOXNUMBER", "box_no"),
    ("BOXNO", "box_no"),
    ("BARCODE", "barcode"),
    ("EAN", "barcode"),
    ("UPC", "barcode"),
    ("STYLE", "item"),
    ("MODEL", "item"),
    ("SKU", "sku"),
    ("BRAND", ""),
    ("CBM", ""),
    ("ARTICLE", "item"),
    ("ITEM", "item"),
    ("QTY", "quantity"),
    ("GW", "gw"),
    ("NW", "nw"),
)


def _pdf_vector_label_info(text: str) -> tuple[bool, str, bool]:
    """Return (label-only, field key, starts with a known table label).

    Some production PDFs draw the complete left label column before drawing the
    complete right value column.  Visitor order is therefore not visual reading
    order, and a single pending-label state can assign every value to the final
    label.  Keep unsupported-but-real slots (BRAND/CBM) with an empty key so they
    still consume their corresponding value during ordered column pairing.
    """

    canonical = re.sub(r"[^A-Z0-9]", "", document_canonical_value(text))
    for prefix, field_key in _PDF_VECTOR_FIELD_PREFIXES:
        if not canonical.startswith(prefix):
            continue
        remainder = canonical[len(prefix):]
        placeholder = re.sub(r"(?:KGS?|LBS?|CMS?|CBM|PCS|CTNS?|OF|X)+", "", remainder)
        label_only = not placeholder
        return label_only, field_key, True
    return False, "", False


def _annotate_pdf_vector_region(lines: list[_PdfVectorDocumentLine], page_height: float) -> None:
    """Attach field context using visual rows, then safe ordered columns."""

    slots = [index for index, line in enumerate(lines) if line.label_only]
    if not slots:
        return

    used_values: set[int] = set()
    row_tolerance = max(3.0, page_height * 0.012)
    # First handle ordinary tables whose label and value share a visual row.
    for slot_index in slots:
        slot = lines[slot_index]
        candidates = [
            index
            for index, value in enumerate(lines)
            if (
                index not in used_values
                and not value.label_only
                and not value.starts_field_label
                and value.x > slot.x + 8.0
                and abs(value.y - slot.y) <= row_tolerance
            )
        ]
        if not candidates:
            continue
        value_index = min(
            candidates,
            key=lambda index: (abs(lines[index].y - slot.y), lines[index].x - slot.x),
        )
        lines[value_index].field_key = slot.field_key
        used_values.add(value_index)

    # Broken text matrices in a few vendor PDFs make later value y-coordinates
    # unusable, while callback order remains label-column then value-column.  Use
    # that fallback only for an unambiguous 1:1 block of at least three rows.
    remaining_slots = [index for index in slots if not any(
        value_index in used_values
        and abs(lines[value_index].y - lines[index].y) <= row_tolerance
        and lines[value_index].x > lines[index].x + 8.0
        for value_index in used_values
    )]
    if len(remaining_slots) < 3:
        return
    label_xs = sorted(lines[index].x for index in remaining_slots)
    label_x = label_xs[len(label_xs) // 2]
    remaining_values = [
        index
        for index, value in enumerate(lines)
        if (
            index not in used_values
            and not value.label_only
            and not value.starts_field_label
            and value.x > label_x + 8.0
        )
    ]
    if (
        len(remaining_values) != len(remaining_slots)
        or not remaining_values
        or max(lines[index].order for index in remaining_slots)
        >= min(lines[index].order for index in remaining_values)
    ):
        return
    for slot_index, value_index in zip(remaining_slots, remaining_values):
        lines[value_index].field_key = lines[slot_index].field_key


def build_pdf_vector_document_items(
    spans: list[PdfTextSpan],
    *,
    page_width: float,
    page_height: float,
    page_index: int,
) -> list[CartonMarkDocumentTextItem]:
    """Keep the two printed mark bodies on an ordinary landscape PDF page.

    The customer's production PDF often has instructions and an order/item
    heading above the red mark frames.  The actual front and side marks occupy
    the lower-left and lower-right areas, so coordinate scoping removes that
    production-only text while retaining additions inside either mark.
    """

    if page_width <= 0 or page_height <= 0 or page_width / page_height < 1.20:
        return []

    regions = (
        ("正唛", [span for span in spans if span.x <= page_width * 0.40]),
        ("侧唛", [span for span in spans if span.x >= page_width * 0.58]),
    )
    items: list[CartonMarkDocumentTextItem] = []
    for region_name, region_spans in regions:
        region_lines: list[_PdfVectorDocumentLine] = []
        for span_order, span in enumerate(region_spans):
            # Visitor callbacks can contain several logical rows in one text
            # object.  Keep every row so additions such as quantity/carton count
            # and side-mark MADE IN CHINA remain independently auditable.
            if span.y > page_height * 0.57:
                continue
            for raw_line in str(span.text or "").replace("\r", "\n").split("\n"):
                line = _clean_pdf_document_line(raw_line)
                if not line:
                    continue
                compact_line = document_canonical_value(line)
                if (
                    compact_line == "20"
                    and (
                        page_width * 0.32 <= span.x <= page_width * 0.40
                        or span.x >= page_width * 0.82
                    )
                ):
                    # Resin/recycling artwork is commonly emitted as separate
                    # PAP and 20 glyphs at the outer edge of each mark panel.
                    continue
                label_only, prefix_key, starts_field_label = _pdf_vector_label_info(line)
                region_lines.append(_PdfVectorDocumentLine(
                    text=line,
                    x=span.x,
                    y=span.y,
                    order=span_order,
                    field_key=prefix_key if starts_field_label else _document_field_key(line),
                    label_only=label_only,
                    starts_field_label=starts_field_label,
                ))
        _annotate_pdf_vector_region(region_lines, page_height)
        for line_number, line in enumerate(region_lines, start=1):
                items.append(CartonMarkDocumentTextItem(
                    text=line.text,
                    location=f"第 {page_index} 页 · {region_name} · 第 {line_number} 行",
                    field_key=line.field_key,
                ))

    if len(items) < 3:
        return []
    fields = extract_fields("\n".join(item.text for item in items), source="pdf_document_vector")
    if len({field.key for field in fields}) < 2:
        return []
    return items


def _extract_pdf_page_vector_document_items(
    page,
    page_index: int,
) -> list[CartonMarkDocumentTextItem]:
    try:
        page_width = float(page.mediabox.width)
        page_height = float(page.mediabox.height)
        spans: list[PdfTextSpan] = []

        def collect_text(text, _cm, tm, _font_dict, _font_size):
            value = str(text or "")
            if not value.strip():
                return
            try:
                spans.append(PdfTextSpan(text=value, x=float(tm[4]), y=float(tm[5])))
            except (IndexError, TypeError, ValueError):
                return

        page.extract_text(visitor_text=collect_text)
        return build_pdf_vector_document_items(
            spans,
            page_width=page_width,
            page_height=page_height,
            page_index=page_index,
        )
    except Exception:
        return []


def extract_pdf_document_items(content: bytes) -> DocumentExtraction:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception as exc:
        raise CartonMarkDocumentConfigurationError(
            "服务器未配置 PDF 文字读取组件"
        ) from exc

    try:
        reader = PdfReader(BytesIO(content), strict=False)
        if reader.is_encrypted and not reader.decrypt(""):
            raise CartonMarkDocumentError("打印 PDF 已加密，无法核对显示文字")
        page_count = len(reader.pages)
    except CartonMarkDocumentError:
        raise
    except Exception as exc:
        raise CartonMarkDocumentError(f"打印 PDF 无法读取：{exc}") from exc
    if page_count == 0:
        raise CartonMarkDocumentError("打印 PDF 没有页面")
    if page_count > MAX_DOCUMENT_PDF_PAGES:
        raise CartonMarkDocumentError("打印 PDF 页数超过 100 页")

    items: list[CartonMarkDocumentTextItem] = []
    reviews: list[tuple[str, str]] = []
    failed_pages: list[int] = []
    empty_text_pages: list[int] = []
    vector_page_count = 0
    for page_index, page in enumerate(reader.pages, start=1):
        vector_items = _extract_pdf_page_vector_document_items(page, page_index)
        if vector_items:
            items.extend(vector_items)
            vector_page_count += 1
            continue
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            failed_pages.append(page_index)
            reviews.append((f"第 {page_index} 页", f"PDF 文字层读取失败：{exc}"))
            continue
        before_count = len(items)
        _append_pdf_lines(items, text, page_index)
        if len(items) == before_count:
            empty_text_pages.append(page_index)

    if items:
        message = ""
        for page_index in empty_text_pages:
            reviews.append((
                f"第 {page_index} 页",
                "该 PDF 页没有可读取文字层，可能为空白页或图片文字页，请人工确认。",
            ))
        if failed_pages:
            message = f"有 {len(failed_pages)} 页文字层读取失败。"
        if empty_text_pages:
            message = (
                (message + " " if message else "")
                + f"有 {len(empty_text_pages)} 页没有可读取文字层。"
            )
        if vector_page_count:
            vector_message = (
                f"已按文字坐标提取 {vector_page_count} 页箱唛正文，"
                "排除红框外制作页眉与说明。"
            )
            message = f"{vector_message} {message}".strip()
        engine = (
            "pypdf-vector-coordinates"
            if vector_page_count == page_count
            else "pypdf-vector-coordinates+pypdf"
            if vector_page_count
            else "pypdf"
        )
        return DocumentExtraction(
            items=items,
            status=CartonMarkExtractionStatus(
                source="print_pdf",
                ok=not failed_pages and not empty_text_pages,
                engine=engine,
                message=message,
                raw_text=clip_text("\n".join(item.text for item in items)),
            ),
            reviews=reviews,
        )

    ocr = _extract_pdf_document_items_with_ocr(content, page_count)
    return ocr


def _append_pdf_lines(
    items: list[CartonMarkDocumentTextItem],
    text: str,
    page_index: int,
) -> None:
    page_start = len(items)
    line_number = 0
    pending_field_key = ""
    for raw_line in str(text or "").replace("\r", "\n").split("\n"):
        line = _clean_pdf_document_line(raw_line)
        if not document_canonical_value(line):
            continue
        if line_number == 0 and re.fullmatch(r"(\d)\s+\1", line):
            # Some flattened recycling artwork is exposed by pypdf as two
            # identical isolated glyphs before the mark body starts.
            continue
        line_number += 1
        if len(line) > MAX_DOCUMENT_ITEM_CHARS:
            raise CartonMarkDocumentError(f"打印 PDF 第 {page_index} 页存在异常超长文字行")
        if len(items) >= MAX_DOCUMENT_TEXT_ITEMS:
            raise CartonMarkDocumentError("打印 PDF 的文字行过多")
        explicit_field_key = _document_field_key(line)
        field_key = explicit_field_key
        if (
            not field_key
            and pending_field_key
            and _document_is_unlabelled_numeric_or_physical(line)
        ):
            field_key = pending_field_key
            pending_field_key = ""
        label_only, _prefix_key, _starts_field_label = _pdf_vector_label_info(line)
        pending_field_key = explicit_field_key if label_only else ""
        items.append(CartonMarkDocumentTextItem(
            text=line,
            location=f"第 {page_index} 页 · 第 {line_number} 行",
            field_key=field_key,
        ))
    _annotate_reverse_order_pdf_labels(items, page_start)


def _annotate_reverse_order_pdf_labels(
    items: list[CartonMarkDocumentTextItem],
    page_start: int,
) -> None:
    """Pair value rows when a PDF emits values before its label column."""

    page_items = items[page_start:]
    index = 0
    while index < len(page_items):
        label_only, _key, _starts = _pdf_vector_label_info(page_items[index].text)
        if not label_only:
            index += 1
            continue
        label_indexes: list[int] = []
        while index < len(page_items):
            is_label, _key, _starts = _pdf_vector_label_info(page_items[index].text)
            if not is_label:
                break
            label_indexes.append(index)
            index += 1
        if len(label_indexes) < 2:
            continue
        first_label = label_indexes[0]
        previous_boundary = max(
            (
                candidate
                for candidate in range(first_label - 1, -1, -1)
                if page_items[candidate].field_key in {"item", "sku", "description"}
            ),
            default=-1,
        )
        value_indexes = [
            candidate
            for candidate in range(previous_boundary + 1, first_label)
            if (
                not page_items[candidate].field_key
                and _document_is_unlabelled_numeric_or_physical(page_items[candidate].text)
            )
        ]
        if len(value_indexes) < len(label_indexes):
            continue
        value_indexes = value_indexes[-len(label_indexes):]
        for label_index, value_index in zip(label_indexes, value_indexes):
            _label_only, field_key, _starts = _pdf_vector_label_info(
                page_items[label_index].text
            )
            if not field_key:
                continue
            original_index = page_start + value_index
            items[original_index] = CartonMarkDocumentTextItem(
                text=items[original_index].text,
                location=items[original_index].location,
                field_key=field_key,
            )


def _extract_pdf_document_items_with_ocr(content: bytes, page_count: int) -> DocumentExtraction:
    if page_count > 20:
        message = "PDF 没有可读取文字层，且超过 20 页，未执行请求时 OCR。"
        return DocumentExtraction(
            items=[],
            status=CartonMarkExtractionStatus(
                source="print_pdf",
                ok=False,
                engine="pypdf",
                message=message,
                raw_text="",
            ),
            reviews=[("打印 PDF", message)],
        )
    try:
        import pypdfium2 as pdfium  # type: ignore
    except Exception as exc:
        raise CartonMarkDocumentConfigurationError(
            "PDF 没有可读取文字层，服务器未配置 PDF OCR 渲染组件"
        ) from exc

    items: list[CartonMarkDocumentTextItem] = []
    failed_pages: list[int] = []
    engines: set[str] = set()
    try:
        document = pdfium.PdfDocument(content)
        for page_index in range(page_count):
            rendered = document[page_index].render(scale=2).to_pil().convert("RGB")
            text, engine = _ocr_document_page(rendered)
            if not text.strip():
                failed_pages.append(page_index + 1)
                continue
            engines.add(engine)
            _append_pdf_lines(items, text, page_index + 1)
    except CartonMarkDocumentConfigurationError:
        raise
    except Exception as exc:
        message = f"PDF 没有可读取文字层，OCR 失败：{exc}"
        return DocumentExtraction(
            items=[],
            status=CartonMarkExtractionStatus(
                source="print_pdf",
                ok=False,
                engine="pypdfium2+ocr",
                message=message,
                raw_text="",
            ),
            reviews=[("打印 PDF", message)],
        )

    if not items:
        message = "PDF 没有可读取文字层，OCR 也未识别到文字。"
        return DocumentExtraction(
            items=[],
            status=CartonMarkExtractionStatus(
                source="print_pdf",
                ok=False,
                engine="pypdfium2+ocr",
                message=message,
                raw_text="",
            ),
            reviews=[("打印 PDF", message)],
        )
    message = "打印 PDF 没有文字层，已使用 OCR；OCR 结果必须人工复核。"
    if failed_pages:
        message += f" 另有 {len(failed_pages)} 页未识别到文字。"
    return DocumentExtraction(
        items=items,
        status=CartonMarkExtractionStatus(
            source="print_pdf",
            ok=not failed_pages,
            engine="+".join(sorted(engines)) or "pypdfium2+ocr",
            message=message,
            raw_text=clip_text("\n".join(item.text for item in items)),
        ),
        reviews=[("打印 PDF", message)],
    )


def _ocr_document_page(image) -> tuple[str, str]:
    rapidocr_engine = get_rapidocr_engine()
    if rapidocr_engine is not None:
        text = rapidocr_image_to_text(rapidocr_engine, resize_for_ocr(image))
        if text.strip():
            return text, "rapidocr-pp-ocrv6"
    try:
        import pytesseract  # type: ignore
    except Exception as exc:
        if rapidocr_engine is not None:
            return "", "rapidocr-pp-ocrv6"
        raise CartonMarkDocumentConfigurationError(
            "服务器未配置 RapidOCR 或 Tesseract OCR 组件"
        ) from exc
    tesseract_cmd, tesseract_lang = configure_tesseract(pytesseract)
    if not tesseract_cmd:
        if rapidocr_engine is not None:
            return "", "rapidocr-pp-ocrv6"
        raise CartonMarkDocumentConfigurationError(missing_tesseract_message())
    text = tesseract_image_to_string(
        pytesseract,
        resize_for_ocr(image),
        lang=tesseract_lang,
        config="--oem 3 --psm 6 -c preserve_interword_spaces=1",
        timeout=20,
    )
    return text, "pytesseract"


def document_canonical_value(value: str) -> str:
    return "".join(character.value for character in _document_text_characters(str(value), 0))


DOCUMENT_FIELD_COMPATIBILITY_GROUPS = (
    frozenset({"item", "sku"}),
    frozenset({"carton_no", "box_no"}),
)


def _document_field_key(text: str) -> str:
    """Infer a line's explicit field without guessing from a bare value."""

    value = str(text or "").strip()
    if not value:
        return ""
    if re.fullmatch(r"\d(?:\s+\d){7,}", value):
        # Human-readable digits printed beneath barcode bars are often emitted
        # one glyph at a time and have no adjacent BARCODE label.
        return "barcode"
    _label_only, prefixed_key, starts_field_label = _pdf_vector_label_info(value)
    if starts_field_label:
        return prefixed_key
    # Prefer a field whose longest alias starts the line.  This disambiguates
    # compact rows such as ITEM DESCRIPTION without using the numeric value as
    # evidence for a different field.
    canonical = document_canonical_value(value)
    candidates: list[tuple[int, str]] = []
    extra_aliases = (("ITEM DESCRIPTION", "description"), ("PRODUCT DESCRIPTION", "description"))
    for alias, key in extra_aliases:
        alias_key = document_canonical_value(alias)
        if canonical.startswith(alias_key):
            candidates.append((len(alias_key), key))
    for definition in FIELD_DEFINITIONS:
        for alias in definition.aliases:
            alias_key = document_canonical_value(alias)
            if alias_key and canonical.startswith(alias_key):
                candidates.append((len(alias_key), definition.key))
    if candidates:
        return max(candidates)[1]
    return ""


def _document_fields_compatible(expected_key: str, actual_key: str) -> bool:
    expected_key = _document_normalized_field_key(expected_key)
    actual_key = _document_normalized_field_key(actual_key)
    if not expected_key or not actual_key:
        return False
    if expected_key == actual_key:
        return True
    return any(
        expected_key in group and actual_key in group
        for group in DOCUMENT_FIELD_COMPATIBILITY_GROUPS
    )


def _document_normalized_field_key(key: str) -> str:
    key_aliases = {
        "gross_weight": "gw",
        "net_weight": "nw",
    }
    return key_aliases.get(key, key)


def _document_text_characters(text: str, item_index: int) -> list[DocumentCharacter]:
    characters: list[DocumentCharacter] = []
    for original_index, original in enumerate(text):
        normalized = unicodedata.normalize("NFKC", original).translate(
            CANONICAL_DOCUMENT_PUNCTUATION
        ).casefold().upper()
        for value in normalized:
            category = unicodedata.category(value)
            if (
                value.isspace()
                or category in {"Cf", "Cc"}
                or value in DOCUMENT_LAYOUT_SEPARATORS
                or value == "_"
            ):
                continue
            characters.append(DocumentCharacter(
                value=value,
                item_index=item_index,
                original_index=original_index,
                break_before=False,
            ))
    return characters


def compare_document_text_items(
    expected_items: list[CartonMarkDocumentTextItem],
    actual_items: list[CartonMarkDocumentTextItem],
) -> list[CartonMarkDocumentComparisonItem]:
    actual_characters: list[DocumentCharacter] = []
    previous_page = ""
    for item_index, item in enumerate(actual_items):
        page = item.location.split(" · ", 1)[0]
        if previous_page and page != previous_page:
            actual_characters.append(DocumentCharacter("\x00", -1, -1, True))
        actual_characters.extend(_document_text_characters(item.text, item_index))
        previous_page = page
    actual_text = "".join(character.value for character in actual_characters)
    used = [character.item_index < 0 for character in actual_characters]
    matches: dict[int, tuple[int, int]] = {}
    canonical_expected = [document_canonical_value(item.text) for item in expected_items]
    exact_search_steps = 0

    for expected_index in sorted(
        range(len(expected_items)),
        key=lambda index: (-len(canonical_expected[index]), index),
    ):
        needle = canonical_expected[expected_index]
        if not needle:
            continue
        start = actual_text.find(needle)
        while start >= 0:
            exact_search_steps += 1
            if exact_search_steps > MAX_DOCUMENT_EXACT_SEARCH_STEPS:
                return [CartonMarkDocumentComparisonItem(
                    status="review",
                    note="文档中存在过多重复文字组合，无法在单次请求内可靠完成逐项配对，请人工复核。",
                )]
            end = start + len(needle)
            if (
                not any(used[start:end])
                and _document_exact_match_is_safe(
                    expected_items[expected_index],
                    actual_items,
                    actual_characters[start:end],
                )
            ):
                matches[expected_index] = (start, end)
                for position in range(start, end):
                    used[position] = True
                break
            start = actual_text.find(needle, start + 1)

    _mark_layout_separators_used(actual_characters, used)
    _mark_repeated_actual_items_used(actual_items, actual_characters, used)
    residuals, partial_residuals = _document_residual_items(
        actual_items,
        actual_characters,
        used,
    )
    ignored_residuals = _document_ignored_residual_indexes(
        residuals,
        expected_items,
        partial_residuals,
        actual_items,
    )
    unmatched_expected = [index for index in range(len(expected_items)) if index not in matches]
    changed_matches: dict[int, int] = {}
    used_residuals: set[int] = set()
    fuzzy_checks = 0
    fuzzy_exhausted = False
    for expected_index in unmatched_expected:
        expected_value = canonical_expected[expected_index]
        best_index = -1
        best_score = 0.0
        for residual_index, residual in enumerate(residuals):
            if residual_index in used_residuals or residual_index in ignored_residuals:
                continue
            actual_value = document_canonical_value(residual.text)
            if not actual_value:
                continue
            if not _document_fuzzy_change_is_safe(
                expected_items[expected_index].text,
                residual.text,
                expected_field_key=expected_items[expected_index].field_key,
                actual_field_key=residual.field_key,
            ):
                continue
            fuzzy_checks += 1
            if fuzzy_checks > MAX_DOCUMENT_FUZZY_CHECKS:
                fuzzy_exhausted = True
                break
            score = _document_changed_similarity(expected_value, actual_value)
            if score > best_score:
                best_score = score
                best_index = residual_index
        if best_index >= 0 and best_score >= 0.62:
            changed_matches[expected_index] = best_index
            used_residuals.add(best_index)
        if fuzzy_exhausted:
            break

    if fuzzy_exhausted:
        return [CartonMarkDocumentComparisonItem(
            status="review",
            note="未配对文字过多，无法在单次请求内可靠区分改字、缺字与新增文字，请人工复核。",
        )]

    comparisons: list[CartonMarkDocumentComparisonItem] = []
    for expected_index, expected in enumerate(expected_items):
        if expected_index in matches:
            start, end = matches[expected_index]
            actual_text_value, actual_location = _document_span_value(
                actual_items,
                actual_characters[start:end],
            )
            comparisons.append(CartonMarkDocumentComparisonItem(
                status="pass",
                expected=expected.text,
                actual=actual_text_value,
                expected_location=expected.location,
                actual_location=actual_location,
            ))
        elif expected_index in changed_matches:
            residual = residuals[changed_matches[expected_index]]
            comparisons.append(CartonMarkDocumentComparisonItem(
                status="changed",
                expected=expected.text,
                actual=residual.text,
                expected_location=expected.location,
                actual_location=residual.location,
                note="Excel 原文在打印 PDF 中存在近似文字，但内容已发生变化。",
            ))
        else:
            comparisons.append(CartonMarkDocumentComparisonItem(
                status="missing",
                expected=expected.text,
                expected_location=expected.location,
                note="Excel 的该段显示文字未在打印 PDF 中找到。",
            ))

    for residual_index, residual in enumerate(residuals):
        if residual_index in used_residuals or residual_index in ignored_residuals:
            continue
        comparisons.append(CartonMarkDocumentComparisonItem(
            status="unexpected",
            actual=residual.text,
            actual_location=residual.location,
            note="打印 PDF 中出现了客户 Excel 没有的文字。",
        ))
    return comparisons


def _document_exact_match_is_safe(
    expected_item: CartonMarkDocumentTextItem,
    actual_items: list[CartonMarkDocumentTextItem],
    characters: list[DocumentCharacter],
) -> bool:
    expected_text = expected_item.text
    expected = document_canonical_value(expected_text)
    item_indexes = {
        character.item_index
        for character in characters
        if character.item_index >= 0
    }
    actual_field_keys = {
        _document_match_field_key(
            actual_items[item_index],
            [character for character in characters if character.item_index == item_index],
        )
        for item_index in item_indexes
    }
    actual_field_keys.discard("")
    expected_field_key = expected_item.field_key
    expected_label_only, expected_label_key, starts_expected_label = (
        _pdf_vector_label_info(expected_text)
    )
    if expected_label_only and starts_expected_label and expected_label_key:
        if any(
            _pdf_vector_label_info(actual_items[item_index].text)[1]
            == expected_label_key
            and _pdf_vector_label_info(actual_items[item_index].text)[2]
            for item_index in item_indexes
        ):
            # A source cell containing only a label can be mis-keyed by a
            # neighbouring Excel header.  The printed label text itself is the
            # stronger evidence; consume only those label characters so a PDF
            # value appended to the same line remains independently auditable.
            return True
    if expected_field_key and any(
        not _document_fields_compatible(expected_field_key, actual_key)
        for actual_key in actual_field_keys
    ):
        return False
    if (
        expected_field_key
        and not actual_field_keys
        and _document_is_unlabelled_numeric_or_physical(expected_text)
    ):
        expected_key = _document_normalized_field_key(expected_field_key)
        if not (
            expected_key in IDENTIFIER_FIELD_KEYS
            and len(expected) >= 6
            and any(
                _document_match_is_slash_delimited_identifier(
                    actual_items[item_index],
                    [
                        character
                        for character in characters
                        if character.item_index == item_index
                    ],
                    expected,
                )
                for item_index in item_indexes
            )
        ):
            return False
    if (
        not expected_field_key
        and actual_field_keys
        and _document_is_unlabelled_numeric_or_physical(expected_text)
    ):
        return False
    if (
        not expected_field_key
        and not actual_field_keys
        and expected.isdigit()
        and len(expected) <= 3
    ):
        return False

    if not expected.isdigit() or len(expected) > 3:
        return True
    identifier_label = re.compile(
        r"\b(?:ITEM|ARTICLE|STYLE|SKU|BAR\s*CODE|BARCODE|PO|P\.O\.?|ORDER|PI)"
        r"\s*(?:NO|NUMBER|#)?\b",
        re.IGNORECASE,
    )
    for item_index in item_indexes:
        actual_text = actual_items[item_index].text
        if not identifier_label.search(actual_text):
            continue
        identifier_values = {
            document_canonical_value(field.value)
            for field in extract_fields(actual_text, source="document_exact_context")
            if field.key in IDENTIFIER_FIELD_KEYS
        }
        if expected not in identifier_values:
            return False
    return True


def _document_match_field_key(
    item: CartonMarkDocumentTextItem,
    characters: list[DocumentCharacter],
) -> str:
    """Resolve field context at the matched substring, not just the whole line."""

    if not characters:
        return item.field_key or _document_field_key(item.text)
    match_start = min(character.original_index for character in characters)
    match_end = max(character.original_index for character in characters) + 1
    matched_text = re.sub(
        r"[^A-Z0-9]",
        "",
        document_canonical_value(item.text[match_start:match_end]),
    )
    matched_candidates = [
        (len(prefix), field_key)
        for prefix, field_key in _PDF_VECTOR_FIELD_PREFIXES
        if field_key and matched_text.startswith(prefix)
    ]
    if matched_candidates:
        return max(matched_candidates)[1]
    before_match = re.sub(
        r"[^A-Z0-9]",
        "",
        document_canonical_value(item.text[:match_start]),
    )
    candidates: list[tuple[int, int, str]] = []
    for prefix, field_key in _PDF_VECTOR_FIELD_PREFIXES:
        if not field_key:
            continue
        position = before_match.rfind(prefix)
        if position >= 0:
            candidates.append((position, len(prefix), field_key))
    if candidates:
        return max(candidates)[2]
    return item.field_key or _document_field_key(item.text)


def _document_match_is_slash_delimited_identifier(
    item: CartonMarkDocumentTextItem,
    characters: list[DocumentCharacter],
    expected: str,
) -> bool:
    if not characters or item.text.count("/") < 2:
        return False
    components = [component.strip() for component in item.text.split("/")]
    if len(components) < 3 or any(not component or len(component) > 40 for component in components):
        return False
    return any(document_canonical_value(component) == expected for component in components)


@lru_cache(maxsize=1)
def _document_field_label_canonicals() -> tuple[str, ...]:
    values = {
        document_canonical_value(value)
        for definition in FIELD_DEFINITIONS
        for value in (definition.label, *definition.aliases)
    }
    values.update({
        document_canonical_value(value)
        for value in (
            "NAME OF SUPPLIER",
            "SUPPLIER NAME",
            "COUNTRY OF ORIGIN",
            "CARTON SIZE(CM)",
            "CARTON SIZE",
            "ITEM DESCRIPTION",
            "PRODUCT DESCRIPTION",
        )
    })
    values.discard("")
    return tuple(sorted(values, key=len, reverse=True))


def _document_ignored_residual_indexes(
    residuals: list[CartonMarkDocumentTextItem],
    expected_items: list[CartonMarkDocumentTextItem],
    partial_residuals: list[bool],
    actual_items: list[CartonMarkDocumentTextItem],
) -> set[int]:
    expected_signatures = {
        document_canonical_value(item.text)
        for item in expected_items
        if document_canonical_value(item.text)
    }
    expected_numbers_by_field: dict[str, set[str]] = {}
    for item in expected_items:
        field_key = _document_normalized_field_key(item.field_key)
        if not field_key:
            continue
        expected_numbers_by_field.setdefault(field_key, set()).update(
            _document_numeric_tokens(item.text)
        )
    actual_by_location = {item.location: item for item in actual_items}
    partial_panel_signatures = {
        (
            residual.location.split(" · ", 1)[0],
            document_canonical_value(actual_by_location[residual.location].text),
        )
        for index, residual in enumerate(residuals)
        if (
            index < len(partial_residuals)
            and partial_residuals[index]
            and residual.location in actual_by_location
        )
    }

    ignored: set[int] = set()
    for index, residual in enumerate(residuals):
        signature = document_canonical_value(residual.text)
        if not signature:
            ignored.add(index)
            continue
        if (
            not (partial_residuals[index] if index < len(partial_residuals) else False)
            and (
                residual.location.split(" · ", 1)[0],
                signature,
            ) in partial_panel_signatures
        ):
            # The same complete PDF line is printed on another panel of the
            # same page.  One occurrence was already partially consumed by an
            # Excel match, so this is an explicit repeated panel—not an
            # independent substring addition.
            ignored.add(index)
            continue
        if _document_residual_is_layout_only(
            residual.text,
            expected_signatures,
            expected_numbers_by_field,
            actual_field_key=residual.field_key or _document_field_key(residual.text),
            is_partial=(
                partial_residuals[index]
                if index < len(partial_residuals)
                else False
            ),
        ):
            ignored.add(index)
    return ignored


def _document_residual_is_layout_only(
    text: str,
    expected_signatures: set[str],
    expected_numbers_by_field: dict[str, set[str]],
    *,
    actual_field_key: str = "",
    is_partial: bool = False,
) -> bool:
    remaining = document_canonical_value(text)
    if not remaining:
        return True
    if not any(character.isalnum() for character in str(text or "")):
        return True

    # A duplicated side panel may be split differently from the front panel.
    # A sufficiently long fragment that is already contained in one Excel value
    # is still the same source wording, not a PDF addition.
    if is_partial and len(remaining) >= 8 and any(
        remaining in expected
        for expected in expected_signatures
        if len(expected) >= len(remaining)
    ):
        return True

    # Remove exact source wording first.  This tolerates the same ITEM, supplier
    # or static origin line being repeated on another panel while still leaving
    # any newly added suffix/prefix behind for reporting.
    cover_values = sorted(
        (
            value
            for value in expected_signatures
            if len(value) >= 3 and (not value.isdigit() or len(value) >= 4)
        ),
        key=len,
        reverse=True,
    )
    if is_partial:
        for value in cover_values:
            remaining = remaining.replace(value, "")
    if not remaining:
        return True

    before_labels = remaining
    remaining = _strip_document_layout_tokens(
        remaining,
        _document_field_label_canonicals(),
    )
    had_explicit_label = remaining != before_labels
    if not remaining:
        return True
    if is_partial and len(remaining) >= 8 and any(
        remaining in expected
        for expected in expected_signatures
        if len(expected) >= len(remaining)
    ):
        return True
    if (
        is_partial
        and remaining.isdigit()
        and len(remaining) <= 2
        and remaining != before_labels
        and document_canonical_value(text).startswith(("DESCRIPTION", "PRODUCTDESCRIPTION"))
        and any(
            expected.startswith(remaining) and len(expected) >= 8
            for expected in expected_signatures
        )
    ):
        return True

    physical_context = bool(re.search(
        r"(?:\b(?:KG|KGS|CM|CBM|PCS|CTNS?)\b|[X×])",
        str(text or "").upper(),
    ))
    if physical_context:
        numbers = _document_numeric_tokens(remaining)
        if not numbers and had_explicit_label:
            placeholder = _strip_document_layout_tokens(
                remaining,
                ("KGS", "KG", "CTNS", "CTN", "PCS", "CBM", "CM", "OF", "X"),
            )
            if not placeholder:
                return True
        expected_numbers = expected_numbers_by_field.get(
            _document_normalized_field_key(actual_field_key),
            set(),
        )
        if numbers and all(number in expected_numbers for number in numbers):
            non_numeric = re.sub(r"[0-9.,+\-]", "", remaining)
            non_numeric = _strip_document_layout_tokens(
                non_numeric,
                ("KGS", "KG", "CTNS", "CTN", "PCS", "CBM", "CM", "OF", "X"),
            )
            if not non_numeric:
                return True
    return False


def _strip_document_layout_tokens(value: str, tokens: tuple[str, ...]) -> str:
    remaining = value
    changed = True
    while remaining and changed:
        changed = False
        for token in tokens:
            if not token:
                continue
            if remaining == token:
                return ""
            # Very short aliases such as PO/X/OF may only consume a whole edge;
            # longer labels can safely be removed when attached to their value.
            if len(token) >= 3 and remaining.startswith(token):
                remaining = remaining[len(token):]
                changed = True
                break
            if len(token) >= 3 and remaining.endswith(token):
                remaining = remaining[:-len(token)]
                changed = True
                break

        if changed:
            continue
        for token in ("KGS", "KG", "CTNS", "CTN", "PCS", "CBM", "CM", "OF", "X"):
            if remaining == token:
                return ""
            if remaining.startswith(token):
                remaining = remaining[len(token):]
                changed = True
                break
            if remaining.endswith(token):
                remaining = remaining[:-len(token)]
                changed = True
                break
    return remaining


def _document_numeric_tokens(value: str) -> set[str]:
    tokens: set[str] = set()
    for match in re.findall(r"[-+]?\d+(?:[.,]\d+)?", str(value or "")):
        try:
            number = Decimal(match.replace(",", "."))
        except InvalidOperation:
            continue
        normalized = format(number.normalize(), "f")
        tokens.add("0" if normalized in {"-0", "+0"} else normalized)
    return tokens


def _document_fuzzy_change_is_safe(
    expected: str,
    actual: str,
    *,
    expected_field_key: str = "",
    actual_field_key: str = "",
) -> bool:
    """Reserve fuzzy changes for wording/labelled identifiers, never bare values.

    A short weight such as 3.12 must not be paired with an unrelated residual 12.
    Without field context the safe result is missing + unexpected, not a confident
    claim that the customer text was changed.
    """

    if (
        expected_field_key
        and actual_field_key
        and not _document_fields_compatible(expected_field_key, actual_field_key)
    ):
        return False
    if (
        not expected_field_key
        and actual_field_key
        and _document_is_unlabelled_numeric_or_physical(expected)
    ):
        return False
    return not (
        _document_is_unlabelled_numeric_or_physical(expected)
        or _document_is_unlabelled_numeric_or_physical(actual)
    )


def _document_is_unlabelled_numeric_or_physical(value: str) -> bool:
    text = str(value or "").upper()
    if not re.search(r"\d", text):
        return False
    without_units = re.sub(
        r"(?:(?<=\d)|\b)(?:KGS?|LBS?|CM|CMS|CBM|PCS|CTNS?|OF)\b|[X×]",
        "",
        text,
    )
    return not bool(re.search(r"[A-WYZ一-龥]", without_units))


def _document_residual_items(
    actual_items: list[CartonMarkDocumentTextItem],
    characters: list[DocumentCharacter],
    used: list[bool],
) -> tuple[list[CartonMarkDocumentTextItem], list[bool]]:
    by_item: dict[int, list[tuple[int, DocumentCharacter]]] = {}
    for position, character in enumerate(characters):
        if character.item_index < 0:
            continue
        by_item.setdefault(character.item_index, []).append((position, character))

    residuals: list[CartonMarkDocumentTextItem] = []
    partial_residuals: list[bool] = []
    for item_index, entries in by_item.items():
        has_used = any(used[position] for position, _character in entries)
        has_unused = any(not used[position] for position, _character in entries)
        is_partial = has_used and has_unused
        group: list[DocumentCharacter] = []
        for position, character in entries:
            if not used[position]:
                group.append(character)
                continue
            if character.value in DOCUMENT_LAYOUT_SEPARATORS:
                continue
            if group:
                _append_document_residual(
                    residuals,
                    partial_residuals,
                    actual_items[item_index],
                    group,
                    is_partial=is_partial,
                )
                group = []
        if group:
            _append_document_residual(
                residuals,
                partial_residuals,
                actual_items[item_index],
                group,
                is_partial=is_partial,
            )
    return residuals, partial_residuals


def _mark_layout_separators_used(
    characters: list[DocumentCharacter],
    used: list[bool],
) -> None:
    for position, character in enumerate(characters):
        if not used[position] and character.value in DOCUMENT_LAYOUT_SEPARATORS:
            used[position] = True


def _mark_repeated_actual_items_used(
    actual_items: list[CartonMarkDocumentTextItem],
    characters: list[DocumentCharacter],
    used: list[bool],
) -> None:
    positions_by_item: dict[int, list[int]] = {}
    for position, character in enumerate(characters):
        if character.item_index >= 0:
            positions_by_item.setdefault(character.item_index, []).append(position)

    fully_used_by_signature: dict[str, list[int]] = {}
    for item_index, positions in positions_by_item.items():
        if not positions or not all(used[position] for position in positions):
            continue
        signature = document_canonical_value(actual_items[item_index].text)
        if signature:
            fully_used_by_signature.setdefault(signature, []).append(item_index)
    if not fully_used_by_signature:
        return

    for item_index, positions in positions_by_item.items():
        signature = document_canonical_value(actual_items[item_index].text)
        source_indexes = fully_used_by_signature.get(signature, [])
        if any(
            (
                _document_locations_are_repeated_panels(
                    actual_items[source_index].location,
                    actual_items[item_index].location,
                )
                or (
                    actual_items[source_index].location.split(" · ", 1)[0]
                    != actual_items[item_index].location.split(" · ", 1)[0]
                    and _document_cross_page_repetition_is_safe(actual_items[item_index])
                )
            )
            for source_index in source_indexes
            if source_index != item_index
        ):
            for position in positions:
                used[position] = True


def _document_locations_are_repeated_panels(left: str, right: str) -> bool:
    left_parts = [part.strip() for part in str(left or "").split(" · ")]
    right_parts = [part.strip() for part in str(right or "").split(" · ")]
    if not left_parts or not right_parts or left_parts[0] != right_parts[0]:
        return False
    panel_names = {"正唛", "侧唛"}
    left_panel = next((part for part in left_parts if part in panel_names), "")
    right_panel = next((part for part in right_parts if part in panel_names), "")
    return bool(left_panel and right_panel and left_panel != right_panel)


def _document_cross_page_repetition_is_safe(item: CartonMarkDocumentTextItem) -> bool:
    canonical = document_canonical_value(item.text)
    if not canonical:
        return False
    if not _document_numeric_tokens(item.text):
        return True
    return (
        (item.field_key or _document_field_key(item.text)) == "description"
        and len(canonical) >= 20
    )


def _append_document_residual(
    residuals: list[CartonMarkDocumentTextItem],
    partial_residuals: list[bool],
    item: CartonMarkDocumentTextItem,
    characters: list[DocumentCharacter],
    *,
    is_partial: bool,
) -> None:
    start = min(character.original_index for character in characters)
    end = max(character.original_index for character in characters) + 1
    text = item.text[start:end].strip(" \t|")
    if document_canonical_value(text):
        residuals.append(CartonMarkDocumentTextItem(
            text=text,
            location=item.location,
            field_key=item.field_key or _document_field_key(item.text),
        ))
        partial_residuals.append(is_partial)


def _document_span_value(
    actual_items: list[CartonMarkDocumentTextItem],
    characters: list[DocumentCharacter],
) -> tuple[str, str]:
    ranges: dict[int, tuple[int, int]] = {}
    for character in characters:
        if character.item_index < 0:
            continue
        current = ranges.get(character.item_index)
        if current is None:
            ranges[character.item_index] = (character.original_index, character.original_index + 1)
        else:
            ranges[character.item_index] = (
                min(current[0], character.original_index),
                max(current[1], character.original_index + 1),
            )
    values: list[str] = []
    locations: list[str] = []
    for item_index, (start, end) in ranges.items():
        values.append(actual_items[item_index].text[start:end])
        locations.append(actual_items[item_index].location)
    return " ".join(values).strip(), " → ".join(locations[:4])


def _document_changed_similarity(expected: str, actual: str) -> float:
    if not expected or not actual:
        return 0.0
    shortest = min(len(expected), len(actual))
    longest = max(len(expected), len(actual))
    if shortest / longest < 0.45:
        return 0.0
    ratio = SequenceMatcher(None, expected, actual, autojunk=False).ratio()
    if expected.isdigit() and actual.isdigit() and len(expected) == len(actual):
        distance = sum(left != right for left, right in zip(expected, actual))
        if distance <= max(1, len(expected) // 5):
            ratio = max(ratio, 0.72)
    return ratio


def summarize_document_comparisons(
    comparisons: list[CartonMarkDocumentComparisonItem],
    *,
    force_review: bool = False,
) -> CartonMarkDocumentCheckSummary:
    pass_count = sum(item.status == "pass" for item in comparisons)
    changed_count = sum(item.status == "changed" for item in comparisons)
    missing_count = sum(item.status == "missing" for item in comparisons)
    unexpected_count = sum(item.status == "unexpected" for item in comparisons)
    review_count = sum(item.status == "review" for item in comparisons)
    if force_review:
        overall_status = "需复核"
    elif changed_count or missing_count or unexpected_count:
        overall_status = "发现差异"
    elif review_count:
        overall_status = "需复核"
    else:
        overall_status = "核对通过"
    return CartonMarkDocumentCheckSummary(
        overall_status=overall_status,
        pass_count=pass_count,
        changed_count=changed_count,
        missing_count=missing_count,
        unexpected_count=unexpected_count,
        review_count=review_count,
    )


def _pdf_document_requires_review(extraction: DocumentExtraction) -> bool:
    engine = extraction.status.engine.casefold()
    return (
        bool(extraction.reviews)
        or not extraction.status.ok
        or "ocr" in engine
        or "pytesseract" in engine
    )


@lru_cache(maxsize=1)
def find_tesseract_cmd() -> str:
    env_cmd = os.getenv("TESSERACT_CMD", "").strip()
    if env_cmd and Path(env_cmd).exists():
        return env_cmd

    path_cmd = shutil.which("tesseract")
    if path_cmd:
        return path_cmd

    for candidate in TESSERACT_CANDIDATE_PATHS:
        if candidate.exists():
            return str(candidate)

    return ""


@lru_cache(maxsize=8)
def get_tesseract_languages(tesseract_cmd: str) -> tuple[str, ...]:
    try:
        completed = subprocess.run(
            [tesseract_cmd, "--list-langs"],
            capture_output=True,
            text=True,
            timeout=10,
            check=False,
        )
    except Exception:
        return ()

    languages = []
    for line in f"{completed.stdout}\n{completed.stderr}".splitlines():
        value = line.strip()
        if not value or value.startswith("List of available"):
            continue
        languages.append(value)

    return tuple(languages)


def select_tesseract_language(tesseract_cmd: str) -> str:
    languages = set(get_tesseract_languages(tesseract_cmd))
    if "eng" in languages and "chi_sim" in languages:
        return "eng+chi_sim"
    if "eng" in languages:
        return "eng"

    for language in languages:
        if language != "osd":
            return language

    return "eng"


def configure_tesseract(pytesseract_module) -> tuple[str, str]:
    tesseract_cmd = find_tesseract_cmd()
    if not tesseract_cmd:
        return "", ""

    pytesseract_module.pytesseract.tesseract_cmd = tesseract_cmd
    return tesseract_cmd, select_tesseract_language(tesseract_cmd)


def missing_tesseract_message() -> str:
    checked_paths = "、".join(str(path) for path in TESSERACT_CANDIDATE_PATHS)
    return f"未找到 tesseract.exe。已检查 PATH、TESSERACT_CMD 和 {checked_paths}。"


def build_carton_mark_auto_check(
    *,
    pdf_bytes: bytes,
    front_image_bytes: bytes,
    side_image_bytes: bytes,
    customer_name: str = "",
    po: str = "",
    item: str = "",
) -> CartonMarkAutoCheckResponse:
    pdf_text, pdf_status = extract_pdf_text(pdf_bytes)
    front_text, front_status = extract_image_text(front_image_bytes, source="front_photo")
    side_text, side_status = extract_image_text(side_image_bytes, source="side_photo")

    metadata = {
        "customer_name": customer_name,
        "po": po,
        "item": item,
    }
    front_pdf_text, side_pdf_text, pdf_layout_status = extract_pdf_template_side_texts(
        pdf_bytes,
        fallback_text=pdf_text,
        photo_text=f"{front_text}\n{side_text}",
        metadata=metadata,
    )
    selected_pdf_text = f"{front_pdf_text}\n{side_pdf_text}"
    template_fields = merge_metadata_fields(
        extract_fields(selected_pdf_text, source="pdf_template"),
        metadata,
        source="template_metadata",
    )
    front_template_fields = merge_metadata_fields(
        extract_fields(front_pdf_text, source="pdf_front_mark"),
        metadata,
        source="template_metadata",
    )
    side_template_fields = merge_metadata_fields(
        extract_fields(side_pdf_text, source="pdf_side_mark"),
        metadata,
        source="template_metadata",
    )

    front_extracted_fields = extract_fields(front_text, source="front_photo")
    side_extracted_fields = extract_fields(side_text, source="side_photo")
    front_expected_fields = enrich_expected_fields_from_actual_values(
        front_template_fields or template_fields,
        front_extracted_fields,
        front_pdf_text,
        source="pdf_front_value_match",
    )
    side_expected_fields = enrich_expected_fields_from_actual_values(
        side_template_fields or template_fields,
        side_extracted_fields,
        side_pdf_text,
        source="pdf_side_value_match",
    )
    front_fields = enrich_fields_from_expected_values(
        front_extracted_fields,
        front_expected_fields,
        front_text,
        source="front_photo_value_match",
    )
    front_fields = enrich_fields_from_expected_values(
        front_fields,
        side_expected_fields,
        front_text,
        source="front_photo_cross_side_value_match",
    )
    side_fields = enrich_fields_from_expected_values(
        side_extracted_fields,
        side_expected_fields,
        side_text,
        source="side_photo_value_match",
    )
    side_fields = enrich_fields_from_expected_values(
        side_fields,
        front_expected_fields,
        side_text,
        source="side_photo_cross_side_value_match",
    )
    comparisons = [
        *compare_observed_photo_fields(
            "front",
            photo_text=front_text,
            photo_status=front_status,
            photo_fields=front_fields,
            front_template_text=front_pdf_text,
            side_template_text=side_pdf_text,
            front_template_fields=front_expected_fields,
            side_template_fields=side_expected_fields,
        ),
        *compare_observed_photo_fields(
            "side",
            photo_text=side_text,
            photo_status=side_status,
            photo_fields=side_fields,
            front_template_text=front_pdf_text,
            side_template_text=side_pdf_text,
            front_template_fields=front_expected_fields,
            side_template_fields=side_expected_fields,
        ),
    ]
    comparisons = cap_comparisons_for_unresolved_pdf_page(comparisons, pdf_layout_status)

    summary = summarize_comparisons(comparisons)
    return CartonMarkAutoCheckResponse(
        summary=summary,
        template_fields=template_fields,
        front_template_fields=front_expected_fields,
        side_template_fields=side_expected_fields,
        front_photo_fields=front_fields,
        side_photo_fields=side_fields,
        comparisons=comparisons,
        extraction=[pdf_status, pdf_layout_status, front_status, side_status],
    )


def build_carton_mark_batch_auto_check(
    *,
    pdf_bytes: bytes,
    front_images: list[tuple[str, bytes]],
    side_images: list[tuple[str, bytes]],
    customer_name: str = "",
    po: str = "",
    item: str = "",
) -> CartonMarkBatchCheckResponse:
    """Check every uploaded front/side image independently against one PDF template."""
    pdf_text, pdf_status = extract_pdf_text(pdf_bytes)
    metadata = {
        "customer_name": customer_name,
        "po": po,
        "item": item,
    }

    def build_item(side: str, image_bytes: bytes) -> CartonMarkAutoCheckResponse:
        source = f"{side}_photo"
        photo_text, photo_status = extract_image_text(image_bytes, source=source)
        front_pdf_text, side_pdf_text, pdf_layout_status = extract_pdf_template_side_texts(
            pdf_bytes,
            fallback_text=pdf_text,
            photo_text=photo_text,
            metadata=metadata,
        )
        template_fields = merge_metadata_fields(
            extract_fields(f"{front_pdf_text}\n{side_pdf_text}", source="pdf_template"),
            metadata,
            source="template_metadata",
        )
        front_template_fields = merge_metadata_fields(
            extract_fields(front_pdf_text, source="pdf_front_mark"),
            metadata,
            source="template_metadata",
        )
        side_template_fields = merge_metadata_fields(
            extract_fields(side_pdf_text, source="pdf_side_mark"),
            metadata,
            source="template_metadata",
        )
        extracted_fields = extract_fields(photo_text, source=source)
        front_expected_fields = enrich_expected_fields_from_actual_values(
            front_template_fields or template_fields,
            extracted_fields,
            front_pdf_text,
            source="pdf_front_value_match",
        )
        side_expected_fields = enrich_expected_fields_from_actual_values(
            side_template_fields or template_fields,
            extracted_fields,
            side_pdf_text,
            source="pdf_side_value_match",
        )
        photo_fields = enrich_fields_from_expected_values(
            extracted_fields,
            front_expected_fields,
            photo_text,
            source=f"{side}_photo_front_value_match",
        )
        photo_fields = enrich_fields_from_expected_values(
            photo_fields,
            side_expected_fields,
            photo_text,
            source=f"{side}_photo_side_value_match",
        )
        comparisons = compare_observed_photo_fields(
            side,
            photo_text=photo_text,
            photo_status=photo_status,
            photo_fields=photo_fields,
            front_template_text=front_pdf_text,
            side_template_text=side_pdf_text,
            front_template_fields=front_expected_fields,
            side_template_fields=side_expected_fields,
        )
        comparisons = cap_comparisons_for_unresolved_pdf_page(comparisons, pdf_layout_status)
        is_front = side == "front"
        return CartonMarkAutoCheckResponse(
            summary=summarize_comparisons(comparisons),
            template_fields=template_fields,
            front_template_fields=front_expected_fields,
            side_template_fields=side_expected_fields,
            front_photo_fields=photo_fields if is_front else [],
            side_photo_fields=photo_fields if not is_front else [],
            comparisons=comparisons,
            extraction=[pdf_status, pdf_layout_status, photo_status],
        )

    items: list[CartonMarkBatchCheckItem] = []
    all_comparisons: list[CartonMarkComparisonItem] = []
    for side, images in (("front", front_images), ("side", side_images)):
        for file_index, (file_name, image_bytes) in enumerate(images):
            result = build_item(side, image_bytes)
            items.append(CartonMarkBatchCheckItem(
                side=side,
                file_name=file_name,
                file_index=file_index,
                result=result,
            ))
            all_comparisons.extend(result.comparisons)

    return CartonMarkBatchCheckResponse(
        summary=summarize_comparisons(all_comparisons),
        items=items,
    )


def extract_pdf_text(pdf_bytes: bytes) -> tuple[str, CartonMarkExtractionStatus]:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        text = fallback_decode_bytes(pdf_bytes)
        return text, CartonMarkExtractionStatus(
            source="pdf_template",
            ok=bool(text.strip()),
            engine="fallback-bytes",
            message="未安装 pypdf，已使用文本字节兜底解析；扫描型 PDF 需要安装 OCR/PDF 解析引擎。",
            raw_text=clip_text(text),
        )

    try:
        reader = PdfReader(BytesIO(pdf_bytes))
        page_texts = [page.extract_text() or "" for page in reader.pages]
        text = "\n".join(page_texts)
        return text, CartonMarkExtractionStatus(
            source="pdf_template",
            ok=bool(text.strip()),
            engine="pypdf",
            message="" if text.strip() else "PDF 没有可抽取文本，可能是扫描图或纯图片 PDF，需要 OCR。",
            raw_text=clip_text(text),
        )
    except Exception as exc:
        text = fallback_decode_bytes(pdf_bytes)
        return text, CartonMarkExtractionStatus(
            source="pdf_template",
            ok=bool(text.strip()),
            engine="fallback-bytes",
            message=f"pypdf 解析失败，已尝试兜底文本解析：{exc}",
            raw_text=clip_text(text),
        )


def extract_pdf_vector_mark_regions(pdf_bytes: bytes) -> list[PdfMarkRegion]:
    """Read repeated mark layouts from vector PDF text before falling back to OCR.

    Customer carton-mark PDFs are commonly exported as one very wide page with
    front / side / front / side marks laid out horizontally.  Raster OCR can
    identify the outer boxes but loses table cells, while the original PDF
    already contains text at precise coordinates.  This path keeps that text
    and splits only deliberately wide pages into their horizontal mark areas.
    """
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        return []

    try:
        reader = PdfReader(BytesIO(pdf_bytes))
        regions: list[PdfMarkRegion] = []

        for page_index, page in enumerate(reader.pages):
            page_width = float(page.mediabox.width)
            page_height = float(page.mediabox.height)
            spans: list[PdfTextSpan] = []

            def collect_text(text, _cm, tm, _font_dict, _font_size):
                value = str(text or "").strip()
                if not value:
                    return
                try:
                    x = float(tm[4])
                    y = float(tm[5])
                except (IndexError, TypeError, ValueError):
                    return
                spans.append(PdfTextSpan(text=value, x=x, y=y))

            plain_text = page.extract_text(visitor_text=collect_text) or ""
            page_regions = classify_pdf_mark_regions(build_pdf_vector_mark_regions(
                spans,
                page_width=page_width,
                page_height=page_height,
                page_index=page_index,
            ))
            regions.extend(page_regions)

            # Most production PDFs are ordinary landscape pages rather than a
            # single extra-wide sheet.  Preserve a layout-mode full-page text
            # region for page anchoring and table values (PI/GW/NW/size).  This
            # avoids raster-OCRing every page merely to find an ITEM number.
            try:
                layout_text = page.extract_text(extraction_mode="layout") or ""
            except (TypeError, ValueError):
                layout_text = plain_text
            full_text = normalize_pdf_layout_text(layout_text or plain_text)
            full_fields = {field.key for field in extract_fields(full_text, source="pdf_page_text")}
            if full_text and full_fields.intersection({"po", "pi_no", "item", "sku", "description", "barcode"}):
                regions.append(PdfMarkRegion(
                    kind="full",
                    text=full_text,
                    box=(0, 0, max(1, int(page_width)), max(1, int(page_height))),
                    page_index=page_index,
                ))

        return sorted(regions, key=pdf_region_position_key)
    except Exception:
        return []


def normalize_pdf_layout_text(text: str) -> str:
    """Collapse pypdf layout padding while repairing split label letters."""
    lines: list[str] = []
    for raw_line in str(text or "").splitlines():
        cells = [cell.strip() for cell in re.split(r"\s{3,}", raw_line.strip()) if cell.strip()]
        if len(cells) >= 2 and len(cells) % 2 == 0:
            middle = len(cells) // 2
            if [normalize_compare_value(cell) for cell in cells[:middle]] == [
                normalize_compare_value(cell) for cell in cells[middle:]
            ]:
                cells = cells[:middle]
        deduped_cells: list[str] = []
        for cell in cells:
            if deduped_cells and normalize_compare_value(cell) == normalize_compare_value(deduped_cells[-1]):
                continue
            deduped_cells.append(cell)
        line = " ".join(deduped_cells)
        line = re.sub(r"\s+", " ", line).strip()
        for _ in range(3):
            line = re.sub(r"\b([A-Za-z])\s+(?=[a-z]{2,}\b)", r"\1", line)
            line = re.sub(r"\b([A-Z])\s+(?=[A-Z]{2,}\b)", r"\1", line)
        line = repair_spaced_field_labels(line)
        cleaned = _clean_pdf_document_line(line)
        if cleaned:
            lines.append(cleaned)
    return "\n".join(lines)


def repair_spaced_field_labels(line: str) -> str:
    repaired = line
    aliases = sorted(
        {alias for definition in FIELD_DEFINITIONS for alias in definition.aliases if len(normalize_alias_key(alias)) >= 5},
        key=lambda value: len(normalize_alias_key(value)),
        reverse=True,
    )
    for alias in aliases:
        parts: list[str] = []
        for character in alias:
            if character.isalnum():
                parts.append(re.escape(character) + r"\s*")
            else:
                parts.append(r"[\s./_\-:]*")
        # pypdf's layout mode can collapse the boundary between a repaired
        # label and its value (``DESCRIPTIONRB`` / ``GRS WT3.18``).  Preserve
        # that boundary explicitly, otherwise the normal alias extractor quite
        # correctly refuses to treat the following value as a label suffix.
        current = repaired

        def replace_label(match: re.Match[str], *, replacement: str = alias) -> str:
            needs_separator = (
                match.end() < len(current)
                and current[match.end()].isalnum()
            )
            return replacement + (" " if needs_separator else "")

        repaired = re.sub(
            "".join(parts),
            replace_label,
            current,
            flags=re.IGNORECASE,
        )
    return repaired


def build_pdf_vector_mark_regions(
    spans: list[PdfTextSpan],
    *,
    page_width: float,
    page_height: float,
    page_index: int = 0,
) -> list[PdfMarkRegion]:
    """Build ordered mark regions from text origins on a deliberately wide page."""
    if page_width <= 0 or page_height <= 0 or page_width / page_height < 2.2:
        return []

    # Decorative footer text is often one long vector string spanning every
    # mark.  It is not part of any mark table and would otherwise bridge the
    # horizontal groups, so only use text in the primary page area.
    usable_spans = [
        span
        for span in spans
        if normalize_compare_value(span.text)
        and 0 <= span.x <= page_width
        and span.y >= page_height * 0.32
    ]
    if len(usable_spans) < 6:
        return []

    group_gap = max(80.0, page_width * 0.12)
    groups: list[list[PdfTextSpan]] = []
    last_x: float | None = None
    for span in sorted(usable_spans, key=lambda item: item.x):
        if last_x is None or span.x - last_x <= group_gap:
            if not groups:
                groups.append([])
            groups[-1].append(span)
        else:
            groups.append([span])
        last_x = span.x

    groups = [group for group in groups if len(group) >= 3]
    if len(groups) < 2:
        return []

    regions: list[PdfMarkRegion] = []
    for group in groups:
        text = build_pdf_vector_region_text(group)
        if not text:
            continue
        left = int(min(span.x for span in group))
        right = int(max(span.x for span in group))
        top = int(max(0, page_height - max(span.y for span in group)))
        bottom = int(max(0, page_height - min(span.y for span in group)))
        regions.append(PdfMarkRegion(
            kind="",
            text=text,
            box=(left, top, right, bottom),
            page_index=page_index,
        ))

    return regions


def build_pdf_vector_region_text(spans: list[PdfTextSpan]) -> str:
    if not spans:
        return ""

    row_threshold = 8.0
    rows: list[list[PdfTextSpan]] = []
    for span in sorted(spans, key=lambda item: (-item.y, item.x)):
        row = next(
            (
                candidate
                for candidate in rows
                if abs(span.y - sum(item.y for item in candidate) / len(candidate)) <= row_threshold
            ),
            None,
        )
        if row is None:
            rows.append([span])
        else:
            row.append(span)

    lines = []
    for row in rows:
        line = " | ".join(item.text for item in sorted(row, key=lambda item: item.x))
        if line.strip():
            lines.append(line)
    return "\n".join(lines)


def extract_pdf_template_side_texts(
    pdf_bytes: bytes,
    *,
    fallback_text: str,
    photo_text: str = "",
    metadata: dict[str, str] | None = None,
) -> tuple[str, str, CartonMarkExtractionStatus]:
    regions, status = extract_pdf_mark_regions(pdf_bytes)
    if not regions:
        return fallback_text, fallback_text, status

    selected_regions, status = select_pdf_mark_page(
        regions,
        photo_text=photo_text,
        metadata=metadata or {},
        extraction_status=status,
    )
    if not selected_regions and status.engine == "pdf-page-selection":
        # Do not populate the UI with fields merged from arbitrary pages.  The
        # photo remains visible, but every result is held for manual review.
        return "", "", status
    front_region = next((region for region in selected_regions if region.kind == "front"), None)
    side_region = next((region for region in selected_regions if region.kind == "side"), None)
    full_text = merge_region_texts(region for region in selected_regions if region.kind == "full")
    front_text = merge_ocr_text_outputs([
        full_text,
        front_region.text if front_region else "",
    ])
    side_text = merge_ocr_text_outputs([
        full_text,
        side_region.text if side_region else "",
    ])
    selected_text = merge_region_texts(selected_regions)

    if not front_text:
        front_text = selected_text or fallback_text
    if not side_text:
        side_text = selected_text or fallback_text

    photo_certification = extract_carton_certification_values(photo_text)
    if photo_certification:
        page_index = (
            (status.matched_page - 1)
            if status.matched_page is not None
            else selected_regions[0].page_index
        )
        certification_text = extract_pdf_certification_stamp_text(pdf_bytes, page_index=page_index)
        if certification_text:
            # The certification seal is normally a separate carton face.  Add
            # it to the side template while keeping the selected page body.
            side_text = merge_ocr_text_outputs([side_text, certification_text])
        else:
            review_reason = "照片包含纸箱认证章，但所选 PDF 页未能可靠提取认证章，请人工复核。"
            status = status.model_copy(update={
                "message": f"{status.message} {review_reason}".strip(),
                "requires_review": True,
                "review_reason": review_reason,
            })

    return front_text, side_text, status


def extract_pdf_certification_stamp_text(pdf_bytes: bytes, *, page_index: int) -> str:
    """OCR only the selected page's certification seal, never every PDF page."""
    try:
        from PIL import ImageEnhance, ImageFilter, ImageOps  # type: ignore
        import pypdfium2 as pdfium  # type: ignore
    except Exception:
        return ""
    engine = get_rapidocr_engine()
    if engine is None:
        return ""
    try:
        try:
            document = pdfium.PdfDocument(pdf_bytes)
        except Exception:
            document = pdfium.PdfDocument(BytesIO(pdf_bytes))
        if page_index < 0 or page_index >= len(document):
            return ""
        image = document[page_index].render(scale=3).to_pil().convert("RGB")
        result = engine(image)
        texts = extract_certification_stamp_ocr_texts(
            engine,
            image,
            result,
            ImageEnhance,
            ImageFilter,
            ImageOps,
        )
        return merge_ocr_text_outputs(texts)
    except Exception:
        return ""


def select_pdf_mark_page(
    regions: list[PdfMarkRegion],
    *,
    photo_text: str,
    metadata: dict[str, str],
    extraction_status: CartonMarkExtractionStatus,
) -> tuple[list[PdfMarkRegion], CartonMarkExtractionStatus]:
    """Choose one PDF page from strong, unique anchors found in the photo.

    A PO repeated on every page is deliberately not useful.  An anchor only
    contributes when it identifies exactly one page; this prevents a partial
    or noisy OCR result from silently comparing the photo with page one.
    """
    pages: dict[int, list[PdfMarkRegion]] = {}
    for region in regions:
        pages.setdefault(region.page_index, []).append(region)
    if not pages:
        return [], extraction_status

    ordered_page_indexes = sorted(pages)
    if len(ordered_page_indexes) == 1:
        page_index = ordered_page_indexes[0]
        selected = select_pdf_page_regions(pages[page_index])
        extraction_note = (
            f" {extraction_status.message}"
            if not extraction_status.ok and extraction_status.message
            else ""
        )
        return selected, CartonMarkExtractionStatus(
            source=extraction_status.source,
            ok=extraction_status.ok,
            engine=extraction_status.engine,
            message=f"PDF 仅 1 页，已选择第 {page_index + 1} 页。{extraction_note}".strip(),
            raw_text=clip_text(merge_region_texts(selected)),
            matched_page=page_index + 1,
            page_count=1,
            match_confidence=1.0,
            requires_review=not extraction_status.ok,
            review_reason=extraction_status.message if not extraction_status.ok else "",
        )

    page_texts = {page_index: merge_region_texts(page_regions) for page_index, page_regions in pages.items()}
    page_fields = {
        page_index: {field.key: field.value for field in extract_fields(text, source="pdf_page_anchor")}
        for page_index, text in page_texts.items()
    }
    photo_fields = {field.key: field.value for field in extract_fields(photo_text, source="photo_page_anchor")}
    page_anchor_values = {
        page_index: extract_strong_anchor_values(text)
        for page_index, text in page_texts.items()
    }
    photo_anchor_values = extract_strong_anchor_values(photo_text)
    for key, value in photo_fields.items():
        if key in {"barcode", "item", "sku", "pi_no", "po", "description"} and value:
            photo_anchor_values.setdefault(key, set()).add(value)

    # Metadata is only a photo anchor when OCR also saw that value.  The caller
    # may hold a stale item number, so metadata alone must never select a page.
    for key in ("barcode", "item", "sku", "pi_no", "po"):
        value = (metadata.get(key) or "").strip()
        if key not in photo_fields and value and photo_text_matches_expected_value(photo_text, value, key):
            photo_fields[key] = value
            photo_anchor_values.setdefault(key, set()).add(value)

    anchor_weights = {
        "barcode": 10,
        "item": 9,
        "sku": 9,
        "pi_no": 8,
        "description": 7,
        "po": 5,
    }
    scores = {page_index: 0 for page_index in ordered_page_indexes}
    reasons: dict[int, list[str]] = {page_index: [] for page_index in ordered_page_indexes}
    for key, weight in anchor_weights.items():
        for actual_value in sorted(photo_anchor_values.get(key, set()), key=len):
            actual_value = actual_value.strip()
            if not actual_value:
                continue
            matching_pages = [
                page_index
                for page_index in ordered_page_indexes
                if any(
                    pdf_page_anchor_matches(key, expected, page_texts[page_index], actual_value)
                    for expected in (
                        page_anchor_values[page_index].get(key, set())
                        or {page_fields[page_index].get(key, "")}
                    )
                )
            ]
            if len(matching_pages) != 1:
                continue
            matched_page = matching_pages[0]
            scores[matched_page] += weight
            reason = f"{FIELD_BY_KEY[key].label}={actual_value}"
            if reason not in reasons[matched_page]:
                reasons[matched_page].append(reason)

    best_score = max(scores.values(), default=0)
    best_pages = [page_index for page_index, score in scores.items() if score == best_score]
    if best_score < 5 or len(best_pages) != 1:
        review_reason = (
            f"PDF 共 {len(ordered_page_indexes)} 页；照片未识别到可唯一定位页面的 ITEM、SKU、条码、PI、描述或 PO 强锚点，"
            "未自动选择页面，请人工复核。"
        )
        return [], CartonMarkExtractionStatus(
            source="pdf_template_regions",
            ok=False,
            engine="pdf-page-selection",
            message=review_reason,
            raw_text=clip_text(photo_text),
            page_count=len(ordered_page_indexes),
            match_confidence=0.0,
            requires_review=True,
            review_reason=review_reason,
        )

    page_index = best_pages[0]
    selected = select_pdf_page_regions(pages[page_index])
    reason = "、".join(reasons[page_index])
    extraction_note = (
        f" {extraction_status.message}"
        if not extraction_status.ok and extraction_status.message
        else ""
    )
    return selected, CartonMarkExtractionStatus(
        source=extraction_status.source,
        ok=extraction_status.ok,
        engine=f"{extraction_status.engine}+photo-anchor",
        message=(
            f"已根据照片强锚点（{reason}）从 {len(ordered_page_indexes)} 页中选择 PDF 第 {page_index + 1} 页。"
            f"{extraction_note}"
        ).strip(),
        raw_text=clip_text(merge_region_texts(selected)),
        matched_page=page_index + 1,
        page_count=len(ordered_page_indexes),
        match_confidence=min(1.0, best_score / 10),
        requires_review=not extraction_status.ok,
        review_reason=extraction_status.message if not extraction_status.ok else "",
    )


def extract_strong_anchor_values(text: str) -> dict[str, set[str]]:
    """Extract standalone identifier tokens even when OCR appends a label/value."""
    anchors: dict[str, set[str]] = {}
    lines = normalize_ocr_lines(text)
    for key in ("barcode", "item", "sku", "pi_no", "po"):
        definition = FIELD_BY_KEY[key]
        values: set[str] = set()
        for line in lines:
            for alias in sorted(definition.aliases, key=lambda value: len(normalize_alias_key(value)), reverse=True):
                match = re.search(
                    rf"(?<![A-Z0-9]){build_flexible_alias_pattern(alias)}(?![A-Z0-9])"
                    rf"\s*(?:(?:N[O0]|NUMBER|NUM)\.?\s*)?[:：#=.\-]*\s*"
                    rf"([A-Z0-9][A-Z0-9./\-]{{2,}})",
                    line,
                    re.IGNORECASE,
                )
                if not match:
                    continue
                token = match.group(1).strip(" .-/")
                if key in {"item", "sku"}:
                    token = re.sub(r"(?:ITEM(?:NO)?|DESC(?:RIPTION)?|CTN|QTY)$", "", token, flags=re.IGNORECASE)
                if token.upper() in {"NUMBER", "MADE", "CHINA", "PIECES", "WAREHOUSE"}:
                    continue
                if key == "barcode" and not barcode_candidate_is_safe(token, line):
                    continue
                if is_searchable_expected_value(key, normalize_compare_value(token)):
                    values.add(token)
        extracted = next((field.value for field in extract_fields(text, source="anchor") if field.key == key), "")
        if extracted and is_searchable_expected_value(key, normalize_compare_value(extracted)):
            values.add(extracted)
        if values:
            anchors[key] = values

    descriptions = [field.value for field in extract_fields(text, source="anchor") if field.key == "description"]
    if descriptions:
        anchors["description"] = set(descriptions)
    return anchors


def select_pdf_page_regions(page_regions: list[PdfMarkRegion]) -> list[PdfMarkRegion]:
    """Keep the primary mark pair plus the selected page's full text layer."""
    full_regions = [region for region in page_regions if region.kind == "full"]
    mark_regions = [region for region in page_regions if region.kind != "full"]
    selected = select_primary_pdf_mark_regions(mark_regions or full_regions)
    seen_text = {normalize_compare_value(region.text) for region in selected}
    for region in sorted(full_regions, key=pdf_region_position_key):
        key = normalize_compare_value(region.text)
        if key and key not in seen_text:
            selected.append(region)
            seen_text.add(key)
    return selected


def pdf_page_anchor_matches(key: str, expected: str, page_text: str, actual: str) -> bool:
    if expected:
        if key == "description":
            expected_key = normalize_compare_value(expected)
            actual_key = normalize_compare_value(actual)
            if min(len(expected_key), len(actual_key)) >= 6 and (
                expected_key in actual_key or actual_key in expected_key
            ):
                return True
        elif key in {"barcode", "item", "sku", "pi_no", "po"}:
            numeric_bias = key in {"barcode", "item", "pi_no", "po"}
            expected_key = normalize_identifier_value(expected, numeric_bias=numeric_bias)
            actual_key = normalize_identifier_value(actual, numeric_bias=numeric_bias)
            return bool(expected_key and expected_key == actual_key)
        elif field_values_match(key, expected, actual):
            return True
    if key in {"barcode", "item", "sku", "pi_no", "po"}:
        return False
    return photo_text_matches_expected_value(page_text, actual, key)


def extract_pdf_mark_regions(pdf_bytes: bytes) -> tuple[list[PdfMarkRegion], CartonMarkExtractionStatus]:
    vector_regions = extract_pdf_vector_mark_regions(pdf_bytes)
    if vector_regions:
        return vector_regions, CartonMarkExtractionStatus(
            source="pdf_template_regions",
            ok=True,
            engine="pypdf-vector-coordinates",
            message="已读取 PDF 全部页面的原始文字坐标，待按照片强锚点选择对应页。",
            raw_text=clip_text(merge_region_texts(vector_regions)),
        )

    try:
        from PIL import ImageFilter  # type: ignore
        import pypdfium2 as pdfium  # type: ignore
        import pytesseract  # type: ignore
    except Exception:
        return [], CartonMarkExtractionStatus(
            source="pdf_template_regions",
            ok=False,
            engine="unconfigured",
            message="PDF 正唛/侧唛区域识别引擎未配置。请安装 pypdfium2、Pillow、pytesseract，并部署 Tesseract 或接入 PaddleOCR。",
            raw_text="",
        )

    try:
        try:
            document = pdfium.PdfDocument(pdf_bytes)
        except Exception:
            document = pdfium.PdfDocument(BytesIO(pdf_bytes))

        if not document:
            raise ValueError("PDF 没有页面")

        regions: list[PdfMarkRegion] = []
        tesseract_cmd, tesseract_lang = configure_tesseract(pytesseract)
        if not tesseract_cmd:
            return [], CartonMarkExtractionStatus(
                source="pdf_template_regions",
                ok=False,
                engine="pypdfium2+pytesseract",
                message=f"{missing_tesseract_message()} 已退回使用 PDF 全文模板字段。",
                raw_text="",
            )

        page_count = min(len(document), MAX_DOCUMENT_PDF_PAGES)
        for page_index in range(page_count):
            rendered = document[page_index].render(scale=3).to_pil().convert("RGB")
            page_regions: list[PdfMarkRegion] = []
            boxes = locate_pdf_mark_boxes(rendered, ImageFilter)
            if len(boxes) < 2:
                boxes = dedupe_boxes([
                    *boxes,
                    *locate_pdf_ocr_text_boxes(rendered, pytesseract, tesseract_lang, ImageFilter),
                ])

            for box in boxes[:8]:
                crop = rendered.crop(box)
                text = pytesseract.image_to_string(crop, lang=tesseract_lang)
                if len(normalize_compare_value(text)) < 4:
                    continue
                page_regions.append(PdfMarkRegion(
                    kind="",
                    text=text,
                    box=box,
                    page_index=page_index,
                ))

            page_regions = classify_pdf_mark_regions(page_regions)
            if not page_regions:
                full_page_text = pytesseract.image_to_string(rendered, lang=tesseract_lang)
                if full_page_text.strip():
                    full_page_box = (0, 0, rendered.width, rendered.height)
                    page_regions = [
                        PdfMarkRegion(
                            kind="front",
                            text=full_page_text,
                            box=full_page_box,
                            page_index=page_index,
                        ),
                        PdfMarkRegion(
                            kind="side",
                            text=full_page_text,
                            box=full_page_box,
                            page_index=page_index,
                        ),
                    ]
            regions.extend(page_regions)

        combined_text = merge_region_texts(regions)
        if not regions:
            return [], CartonMarkExtractionStatus(
                source="pdf_template_regions",
                ok=False,
                engine="pypdfium2+pytesseract",
                message="未能在 PDF 中定位正唛/侧唛区域，已退回使用 PDF 全文模板字段。",
                raw_text="",
            )

        return regions, CartonMarkExtractionStatus(
            source="pdf_template_regions",
            ok=True,
            engine="pypdfium2+pytesseract",
            message=f"已 OCR 提取 PDF 全部 {page_count} 页，待按照片强锚点选择对应页。",
            raw_text=clip_text(combined_text),
        )
    except Exception as exc:
        return [], CartonMarkExtractionStatus(
            source="pdf_template_regions",
            ok=False,
            engine="pypdfium2+pytesseract",
            message=f"PDF 正唛/侧唛区域识别失败，已退回使用 PDF 全文模板字段：{exc}",
            raw_text="",
        )


def locate_pdf_mark_boxes(image, image_filter) -> list[tuple[int, int, int, int]]:
    max_analysis_width = 1800
    scale = min(1.0, max_analysis_width / max(image.width, 1))
    analysis_image = image
    if scale < 1:
        analysis_image = image.resize((int(image.width * scale), int(image.height * scale)))

    gray = analysis_image.convert("L")
    mask = gray.point(lambda pixel: 255 if pixel < 125 else 0)
    mask = mask.filter(image_filter.MaxFilter(21))
    mask = mask.filter(image_filter.MinFilter(5))
    boxes = find_connected_boxes(mask)

    scaled_boxes = []
    for left, top, right, bottom in boxes:
        scaled_boxes.append(expand_box(
            (
                int(left / scale),
                int(top / scale),
                int(right / scale),
                int(bottom / scale),
            ),
            image.width,
            image.height,
            padding=28,
        ))

    return dedupe_boxes(scaled_boxes)


def locate_pdf_ocr_text_boxes(image, pytesseract_module, tesseract_lang: str, image_filter) -> list[tuple[int, int, int, int]]:
    try:
        from PIL import Image, ImageDraw  # type: ignore
        from pytesseract import Output  # type: ignore
    except Exception:
        return []

    try:
        data = pytesseract_module.image_to_data(
            image,
            lang=tesseract_lang,
            output_type=Output.DICT,
            config="--psm 11",
        )
    except Exception:
        return []

    max_analysis_width = 1800
    scale = min(1.0, max_analysis_width / max(image.width, 1))
    analysis_size = (int(image.width * scale), int(image.height * scale))
    mask = Image.new("L", analysis_size, 0)
    draw = ImageDraw.Draw(mask)
    padding_x = max(10, int(analysis_size[0] * 0.012))
    padding_y = max(8, int(analysis_size[1] * 0.018))

    for index, raw_text in enumerate(data.get("text", [])):
        text = str(raw_text or "").strip()
        if len(normalize_compare_value(text)) < 2:
            continue

        try:
            confidence = float(data.get("conf", [])[index])
        except Exception:
            confidence = 0
        if confidence < 15:
            continue

        left = int(data.get("left", [])[index] * scale)
        top = int(data.get("top", [])[index] * scale)
        width = int(data.get("width", [])[index] * scale)
        height = int(data.get("height", [])[index] * scale)
        if width < 3 or height < 3:
            continue

        draw.rectangle(
            (
                max(0, left - padding_x),
                max(0, top - padding_y),
                min(analysis_size[0], left + width + padding_x),
                min(analysis_size[1], top + height + padding_y),
            ),
            fill=255,
        )

    mask = mask.filter(image_filter.MaxFilter(31))
    mask = mask.filter(image_filter.MinFilter(5))
    boxes = find_connected_boxes(mask)
    boxes = filter_plausible_mark_boxes(boxes, analysis_size[0], analysis_size[1])

    scaled_boxes = []
    for left, top, right, bottom in boxes:
        scaled_boxes.append(expand_box(
            (
                int(left / scale),
                int(top / scale),
                int(right / scale),
                int(bottom / scale),
            ),
            image.width,
            image.height,
            padding=36,
        ))

    return scaled_boxes


def filter_plausible_mark_boxes(
    boxes: list[tuple[int, int, int, int]],
    image_width: int,
    image_height: int,
) -> list[tuple[int, int, int, int]]:
    filtered = []
    min_width = max(45, int(image_width * 0.035))
    min_height = max(45, int(image_height * 0.07))
    max_width = int(image_width * 0.45)
    max_height = int(image_height * 0.55)

    for box in boxes:
        width = box[2] - box[0]
        height = box[3] - box[1]
        aspect_ratio = width / max(height, 1)
        if width < min_width or height < min_height:
            continue
        if width > max_width or height > max_height:
            continue
        if aspect_ratio < 0.35 or aspect_ratio > 3.6:
            continue
        filtered.append(box)

    return filtered


def find_connected_boxes(
    mask,
    *,
    min_width_ratio: float = 0.035,
    min_height_ratio: float = 0.03,
    min_area_ratio: float = 0.0008,
    max_area_ratio: float = 0.18,
) -> list[tuple[int, int, int, int]]:
    width, height = mask.size
    pixels = mask.load()
    visited = bytearray(width * height)
    boxes: list[tuple[int, int, int, int]] = []
    min_width = max(28, int(width * min_width_ratio))
    min_height = max(22, int(height * min_height_ratio))
    min_area = max(500, int(width * height * min_area_ratio))
    max_area = int(width * height * max_area_ratio)

    for y in range(height):
        for x in range(width):
            index = y * width + x
            if visited[index] or pixels[x, y] < 128:
                continue

            stack = [(x, y)]
            visited[index] = 1
            left = right = x
            top = bottom = y
            count = 0

            while stack:
                current_x, current_y = stack.pop()
                count += 1
                left = min(left, current_x)
                right = max(right, current_x)
                top = min(top, current_y)
                bottom = max(bottom, current_y)

                for next_x, next_y in (
                    (current_x + 1, current_y),
                    (current_x - 1, current_y),
                    (current_x, current_y + 1),
                    (current_x, current_y - 1),
                ):
                    if next_x < 0 or next_y < 0 or next_x >= width or next_y >= height:
                        continue
                    next_index = next_y * width + next_x
                    if visited[next_index] or pixels[next_x, next_y] < 128:
                        continue
                    visited[next_index] = 1
                    stack.append((next_x, next_y))

            box_width = right - left + 1
            box_height = bottom - top + 1
            box_area = box_width * box_height
            if box_width < min_width or box_height < min_height:
                continue
            if box_area < min_area or box_area > max_area:
                continue

            boxes.append((left, top, right, bottom))

    return sorted(boxes, key=lambda box: (box[2] - box[0]) * (box[3] - box[1]), reverse=True)


def expand_box(box: tuple[int, int, int, int], image_width: int, image_height: int, *, padding: int) -> tuple[int, int, int, int]:
    left, top, right, bottom = box
    return (
        max(0, left - padding),
        max(0, top - padding),
        min(image_width, right + padding),
        min(image_height, bottom + padding),
    )


def dedupe_boxes(boxes: list[tuple[int, int, int, int]]) -> list[tuple[int, int, int, int]]:
    kept: list[tuple[int, int, int, int]] = []
    for box in boxes:
        if all(box_iou(box, kept_box) < 0.55 for kept_box in kept):
            kept.append(box)

    return kept


def box_iou(left_box: tuple[int, int, int, int], right_box: tuple[int, int, int, int]) -> float:
    left = max(left_box[0], right_box[0])
    top = max(left_box[1], right_box[1])
    right = min(left_box[2], right_box[2])
    bottom = min(left_box[3], right_box[3])
    if right <= left or bottom <= top:
        return 0

    intersection = (right - left) * (bottom - top)
    left_area = (left_box[2] - left_box[0]) * (left_box[3] - left_box[1])
    right_area = (right_box[2] - right_box[0]) * (right_box[3] - right_box[1])
    union = left_area + right_area - intersection
    return intersection / union if union else 0


def classify_pdf_mark_regions(regions: list[PdfMarkRegion]) -> list[PdfMarkRegion]:
    if not regions:
        return []

    widths = [region.box[2] - region.box[0] for region in regions]
    width_threshold = (max(widths) + min(widths)) / 2 if len(widths) > 1 else widths[0]
    classified: list[PdfMarkRegion] = []

    for region in regions:
        width = region.box[2] - region.box[0]
        height = max(region.box[3] - region.box[1], 1)
        aspect_ratio = width / height
        normalized_text = normalize_compare_value(region.text)
        front_score = sum(
            marker in normalized_text
            for marker in ("IMPORTADOR", "DIRECCION", "PROVEEDOR", "RFC")
        )
        side_score = sum(
            marker in normalized_text
            for marker in ("DESCRIPCION", "PIEZASPORBULTO", "CAJANUMERO", "MEAS", "SECCION")
        )
        if front_score != side_score:
            kind = "front" if front_score > side_score else "side"
        else:
            kind = "front" if width >= width_threshold or aspect_ratio >= 1.08 else "side"
        classified.append(PdfMarkRegion(
            kind=kind,
            text=region.text,
            box=region.box,
            page_index=region.page_index,
        ))

    return sorted(classified, key=pdf_region_position_key)


def select_primary_pdf_mark_regions(regions: list[PdfMarkRegion]) -> list[PdfMarkRegion]:
    if not regions:
        return []

    ordered = sorted(regions, key=pdf_region_position_key)
    # A front/side pair must always come from the same PDF page.  Callers that
    # need a later page first narrow ``regions`` with the photo anchor selector.
    first_page_index = ordered[0].page_index
    ordered = [region for region in ordered if region.page_index == first_page_index]
    front_region = next((region for region in ordered if region.kind == "front"), None)
    if front_region is None:
        front_region = ordered[0]

    front_index = next(
        (index for index, region in enumerate(ordered) if region is front_region),
        0,
    )
    side_region = next(
        (region for region in ordered[front_index + 1:] if region.kind == "side"),
        None,
    )
    if side_region is None:
        side_region = next(
            (region for region in ordered if region.kind == "side" and region is not front_region),
            None,
        )
    if side_region is None:
        side_region = next((region for region in ordered[front_index + 1:] if region is not front_region), None)

    selected = [
        PdfMarkRegion(
            kind="front",
            text=front_region.text,
            box=front_region.box,
            page_index=front_region.page_index,
        ),
    ]
    if side_region is not None:
        selected.append(PdfMarkRegion(
            kind="side",
            text=side_region.text,
            box=side_region.box,
            page_index=side_region.page_index,
        ))

    return selected


def pdf_region_position_key(region: PdfMarkRegion) -> tuple[int, int, int, int]:
    left, top, right, bottom = region.box
    return (region.page_index, left, top, -((right - left) * (bottom - top)))


def merge_region_texts(regions) -> str:
    return "\n".join(region.text for region in regions if region.text.strip())


@lru_cache(maxsize=1)
def get_rapidocr_engine():
    try:
        from rapidocr import RapidOCR  # type: ignore

        return RapidOCR()
    except Exception:
        return None


def extract_image_text(image_bytes: bytes, *, source: str) -> tuple[str, CartonMarkExtractionStatus]:
    try:
        from PIL import Image, ImageEnhance, ImageFilter, ImageOps  # type: ignore
    except Exception:
        return "", CartonMarkExtractionStatus(
            source=source,
            ok=False,
            engine="unconfigured",
            message="图片 OCR 引擎未配置。请在后端安装 Pillow 及 RapidOCR 或 Tesseract。",
            raw_text="",
        )

    try:
        image = ImageOps.exif_transpose(Image.open(BytesIO(image_bytes))).convert("RGB")
        variants = build_photo_ocr_variants(image, ImageEnhance, ImageFilter, ImageOps)
        rapidocr_texts = []
        spatial_evidence_lines: list[str] = []
        rapidocr_engine = get_rapidocr_engine()

        if rapidocr_engine is not None:
            # PP-OCR can detect text lines on perspective/low-contrast carton
            # labels.  Test the whole label plus the most relevant crops before
            # falling back to the older character-level Tesseract path.
            for variant_index, variant in enumerate(variants[:8]):
                try:
                    rapidocr_result = rapidocr_engine(variant)
                except Exception:
                    rapidocr_result = None
                text = rapidocr_result_to_text(rapidocr_result)
                if text.strip():
                    rapidocr_texts.append(text)
                if variant_index == 0 and rapidocr_result is not None:
                    spatial_evidence_lines = rapidocr_photo_panel_evidence_lines(
                        rapidocr_result
                    )
                    rapidocr_texts.extend(extract_certification_stamp_ocr_texts(
                        rapidocr_engine,
                        variant,
                        rapidocr_result,
                        ImageEnhance,
                        ImageFilter,
                        ImageOps,
                    ))
                candidate_text = merge_ocr_text_outputs([
                    *rapidocr_texts,
                    "\n".join(spatial_evidence_lines),
                ])
                if len(extract_fields(candidate_text, source=source)) >= 3:
                    break

        rapidocr_text = merge_ocr_text_outputs([
            *rapidocr_texts,
            "\n".join(spatial_evidence_lines),
        ])
        if rapidocr_text and len(extract_fields(rapidocr_text, source=source)) >= 3:
            return rapidocr_text, CartonMarkExtractionStatus(
                source=source,
                ok=True,
                engine="rapidocr-pp-ocrv6",
                message="",
                raw_text=clip_text(rapidocr_text),
            )

        try:
            import pytesseract  # type: ignore
        except Exception:
            if rapidocr_text:
                return rapidocr_text, CartonMarkExtractionStatus(
                    source=source,
                    ok=True,
                    engine="rapidocr-pp-ocrv6",
                    message="RapidOCR 已识别到部分文字；请根据核验结果复核未提取字段。",
                    raw_text=clip_text(rapidocr_text),
                )
            raise RuntimeError("RapidOCR 与 pytesseract 均不可用")

        tesseract_cmd, tesseract_lang = configure_tesseract(pytesseract)
        if not tesseract_cmd:
            if rapidocr_text:
                return rapidocr_text, CartonMarkExtractionStatus(
                    source=source,
                    ok=True,
                    engine="rapidocr-pp-ocrv6",
                    message="RapidOCR 已识别到部分文字；请根据核验结果复核未提取字段。",
                    raw_text=clip_text(rapidocr_text),
                )
            return "", CartonMarkExtractionStatus(
                source=source,
                ok=False,
                engine="rapidocr+pytesseract",
                message=missing_tesseract_message(),
                raw_text="",
            )

        texts = list(rapidocr_texts)
        primary_configs = (
            "--oem 3 --psm 6 -c preserve_interword_spaces=1",
            "--oem 3 --psm 4 -c preserve_interword_spaces=1",
            "--oem 3 --psm 11",
        )
        crop_configs = (
            "--oem 3 --psm 6 -c preserve_interword_spaces=1",
            "--oem 3 --psm 11",
        )
        for index, variant in enumerate(variants):
            configs = primary_configs if index < 3 else crop_configs
            for config in configs:
                text = tesseract_image_to_string(
                    pytesseract,
                    variant,
                    lang=tesseract_lang,
                    config=config,
                    timeout=10,
                )
                if text.strip():
                    texts.append(text)
            if index < 8:
                table_text = tesseract_image_to_table_text(
                    pytesseract,
                    variant,
                    lang=tesseract_lang,
                    config="--oem 3 --psm 6 -c preserve_interword_spaces=1",
                    timeout=10,
                )
                if table_text.strip():
                    texts.append(table_text)

        text = merge_ocr_text_outputs(texts)
        return text, CartonMarkExtractionStatus(
            source=source,
            ok=bool(text.strip()),
            engine="rapidocr-pp-ocrv6+pytesseract-fallback" if rapidocr_engine is not None else "pytesseract-multi-pass",
            message="" if text.strip() else "图片 OCR 已尝试 PP-OCR、整图、箱唛候选区域裁剪和轻微旋转，仍未识别到文字；请镜头正对单块箱唛、让标签占画面约 70%，避开反光后重拍，或先框选箱唛区域再核对。",
            raw_text=clip_text(text),
        )
    except Exception as exc:
        return "", CartonMarkExtractionStatus(
            source=source,
            ok=False,
            engine="rapidocr+pytesseract",
            message=f"图片 OCR 失败：{exc}",
            raw_text="",
        )


def rapidocr_image_to_text(engine, image) -> str:
    try:
        return rapidocr_result_to_text(engine(image))
    except Exception:
        return ""


def rapidocr_result_to_text(result) -> str:
    if result is None:
        return ""

    try:
        texts = tuple(str(value or "").strip() for value in getattr(result, "txts", ()) or ())
    except Exception:
        return ""

    raw_lines = [text for text in texts if text]
    boxes = getattr(result, "boxes", None)
    scores = getattr(result, "scores", None)
    words: list[OcrWord] = []

    if boxes is not None:
        for index, text in enumerate(texts):
            if not text:
                continue
            try:
                points = boxes[index]
                xs = [float(point[0]) for point in points]
                ys = [float(point[1]) for point in points]
                score = float(scores[index]) if scores is not None else 0.8
            except (IndexError, TypeError, ValueError):
                continue
            if not xs or not ys:
                continue
            left = int(min(xs))
            top = int(min(ys))
            words.append(OcrWord(
                text=text,
                left=left,
                top=top,
                right=max(left + 1, int(max(xs))),
                bottom=max(top + 1, int(max(ys))),
                confidence=score * 100,
            ))

    table_text = build_ocr_word_table_text(words)
    lines = []
    if table_text.strip():
        lines.append(table_text)
    lines.extend(raw_lines)
    return merge_ocr_text_outputs(lines)


PHOTO_PANEL_COUNT_RE = re.compile(r"\[PHOTO_PANEL_COUNT:(\d+)\]", re.IGNORECASE)
PHOTO_COUNTRY_PANEL_RE = re.compile(
    r"\[PHOTO_COUNTRY_PANEL:(\d+):([A-Z][A-Z .'-]{1,40})\]",
    re.IGNORECASE,
)


def rapidocr_photo_panel_evidence_lines(result) -> list[str]:
    """Preserve first-pass spatial evidence before OCR variants are merged.

    Re-running crops and rotations is useful for text recovery but duplicates
    lines, so plain-text occurrence counts cannot prove that two carton faces
    were photographed.  The primary detector has stable boxes: repeated strong
    identifiers establish the number of visible mark panels, while an origin
    phrase is counted only at its own detected box.
    """

    try:
        texts = tuple(str(value or "").strip() for value in getattr(result, "txts", ()) or ())
        raw_boxes = getattr(result, "boxes", None)
        boxes = tuple(raw_boxes) if raw_boxes is not None else ()
    except Exception:
        return []
    if not texts or not boxes:
        return []

    anchors: dict[tuple[str, str], list[tuple[float, float]]] = {}
    origins: list[tuple[str, float, float]] = []
    for index, detected_text in enumerate(texts):
        try:
            points = boxes[index]
            xs = [float(point[0]) for point in points]
            ys = [float(point[1]) for point in points]
        except (IndexError, TypeError, ValueError):
            continue
        if not xs or not ys:
            continue
        center = ((min(xs) + max(xs)) / 2, (min(ys) + max(ys)) / 2)
        origin = extract_country_of_origin_value(detected_text)
        if origin:
            origins.append((origin, *center))

        for field in extract_fields(detected_text, source="photo_panel_evidence"):
            if field.key not in {"item", "sku", "barcode"}:
                continue
            value = normalize_identifier_value(field.value, numeric_bias=True)
            if len(value) < 3 or not any(character.isdigit() for character in value):
                continue
            anchors.setdefault((field.key, value), []).append(center)

    panel_count = max((len(centers) for centers in anchors.values()), default=0)
    lines: list[str] = []
    if panel_count:
        lines.append(f"[PHOTO_PANEL_COUNT:{panel_count}]")
    for index, (origin, _x, _y) in enumerate(origins, start=1):
        lines.append(f"[PHOTO_COUNTRY_PANEL:{index}:{origin}]")
    return lines


def extract_certification_stamp_ocr_texts(
    engine,
    image,
    result,
    image_enhance,
    image_filter,
    image_ops,
) -> list[str]:
    """Enlarge a certification stamp located by its isolated CHINA footer."""
    try:
        texts = tuple(str(value or "").strip() for value in getattr(result, "txts", ()) or ())
        raw_boxes = getattr(result, "boxes", None)
        boxes = tuple(raw_boxes) if raw_boxes is not None else ()
    except Exception:
        return []
    if not texts or not boxes:
        return []

    detections: list[tuple[str, float, float, float, float]] = []
    for index, detected_text in enumerate(texts):
        try:
            points = boxes[index]
            xs = [float(point[0]) for point in points]
            ys = [float(point[1]) for point in points]
        except (IndexError, TypeError, ValueError):
            continue
        if xs and ys:
            detections.append((detected_text, min(xs), min(ys), max(xs), max(ys)))

    stamp_texts: list[str] = []
    radius = max(120, int(min(image.width, image.height) * 0.16))
    for detected_text, left, top, right, bottom in detections:
        if normalize_alias_key(detected_text) != "CHINA":
            continue
        center_x = (left + right) / 2
        center_y = (top + bottom) / 2
        nearby_numbers = 0
        for candidate_text, candidate_left, candidate_top, candidate_right, candidate_bottom in detections:
            candidate_x = (candidate_left + candidate_right) / 2
            candidate_y = (candidate_top + candidate_bottom) / 2
            if abs(candidate_x - center_x) > radius * 1.4 or abs(candidate_y - center_y) > radius * 2.2:
                continue
            nearby_numbers += len(re.findall(r"(?<!\d)\d{2,3}(?!\d)", candidate_text))
        if nearby_numbers < 3:
            continue

        crop_box = (
            max(0, int(center_x - radius)),
            max(0, int(center_y - radius * 1.25)),
            min(image.width, int(center_x + radius)),
            min(image.height, int(center_y + radius * 0.35)),
        )
        crop = image.crop(crop_box)
        if (bottom - top) > (right - left) * 1.25:
            crop = crop.rotate(270, expand=True, fillcolor="white")

        variants: list = []
        add_photo_region_variants(
            crop,
            variants,
            set(),
            image_enhance,
            image_filter,
            image_ops,
        )
        for variant in variants[:2]:
            text = rapidocr_image_to_text(engine, variant)
            if extract_carton_certification_values(text):
                stamp_texts.append(text)
        if stamp_texts:
            break
    return stamp_texts


def tesseract_image_to_string(pytesseract_module, image, *, lang: str, config: str, timeout: int) -> str:
    try:
        return pytesseract_module.image_to_string(image, lang=lang, config=config, timeout=timeout)
    except TypeError:
        return pytesseract_module.image_to_string(image, lang=lang, config=config)
    except RuntimeError:
        return ""


def tesseract_image_to_table_text(pytesseract_module, image, *, lang: str, config: str, timeout: int) -> str:
    try:
        from pytesseract import Output  # type: ignore
    except Exception:
        return ""

    try:
        data = pytesseract_module.image_to_data(
            image,
            lang=lang,
            config=config,
            output_type=Output.DICT,
            timeout=timeout,
        )
    except TypeError:
        data = pytesseract_module.image_to_data(
            image,
            lang=lang,
            config=config,
            output_type=Output.DICT,
        )
    except RuntimeError:
        return ""

    return build_ocr_word_table_text(collect_ocr_words(data))


def collect_ocr_words(data: dict) -> list[OcrWord]:
    words: list[OcrWord] = []
    texts = data.get("text", [])
    for index, raw_text in enumerate(texts):
        text = str(raw_text or "").strip()
        if not re.search(r"[A-Z0-9一-龥]", text, re.IGNORECASE):
            continue

        try:
            confidence = float(data.get("conf", [])[index])
        except Exception:
            confidence = 0
        if confidence < 8 and len(normalize_compare_value(text)) < 3:
            continue

        try:
            left = int(float(data.get("left", [])[index]))
            top = int(float(data.get("top", [])[index]))
            width = int(float(data.get("width", [])[index]))
            height = int(float(data.get("height", [])[index]))
        except Exception:
            continue
        if width <= 0 or height <= 0:
            continue

        words.append(OcrWord(
            text=text,
            left=left,
            top=top,
            right=left + width,
            bottom=top + height,
            confidence=confidence,
        ))

    return words


def build_ocr_word_table_text(words: list[OcrWord]) -> str:
    if not words:
        return ""

    rows = group_ocr_words_into_rows(words)
    lines = []
    lines.extend(words_to_spaced_line(row) for row in rows)
    lines.extend(build_label_value_lines_from_rows(rows))
    return merge_ocr_text_outputs(lines)


def group_ocr_words_into_rows(words: list[OcrWord]) -> list[list[OcrWord]]:
    sorted_words = sorted(words, key=lambda word: (word.top + word.bottom, word.left))
    median_height = median_number([word.bottom - word.top for word in sorted_words]) or 12
    row_threshold = max(10, int(median_height * 0.75))
    rows: list[list[OcrWord]] = []

    for word in sorted_words:
        center_y = (word.top + word.bottom) // 2
        best_row = None
        best_distance = row_threshold + 1
        for row in rows:
            row_center = sum((item.top + item.bottom) // 2 for item in row) // len(row)
            distance = abs(center_y - row_center)
            if distance <= row_threshold and distance < best_distance:
                best_row = row
                best_distance = distance

        if best_row is None:
            rows.append([word])
        else:
            best_row.append(word)

    return [sorted(row, key=lambda word: word.left) for row in rows]


def words_to_spaced_line(words: list[OcrWord]) -> str:
    if not words:
        return ""

    sorted_words = sorted(words, key=lambda word: word.left)
    widths = [word.right - word.left for word in sorted_words]
    median_width = median_number(widths) or 16
    parts = [sorted_words[0].text]
    previous = sorted_words[0]

    for word in sorted_words[1:]:
        gap = word.left - previous.right
        parts.append(" | " if gap > median_width * 1.8 else " ")
        parts.append(word.text)
        previous = word

    return "".join(parts)


def build_label_value_lines_from_rows(rows: list[list[OcrWord]]) -> list[str]:
    lines: list[str] = []
    for row in rows:
        if len(row) < 2:
            continue

        row_line = words_to_spaced_line(row)
        for definition in FIELD_DEFINITIONS:
            alias_match = find_alias_word_span(row, definition)
            if alias_match is None:
                continue

            alias_text, alias_right = alias_match
            value_words = [
                word
                for word in row
                if word.left >= alias_right - 2 and not word_looks_like_any_alias(word.text)
            ]
            value_text = clean_field_value(" ".join(word.text for word in value_words))
            if is_plausible_field_value(value_text, definition):
                lines.append(f"{alias_text} {value_text}")
            elif line_looks_like_alias(row_line, alias_text):
                lines.append(row_line)

    return lines


def find_alias_word_span(row: list[OcrWord], definition: FieldDefinition) -> tuple[str, int] | None:
    normalized_words = [normalize_compare_value(word.text) for word in row]
    for alias in sorted(definition.aliases, key=lambda value: len(normalize_compare_value(value)), reverse=True):
        alias_key = normalize_compare_value(alias)
        if not alias_key:
            continue

        for start in range(len(row)):
            combined = ""
            for end in range(start, min(len(row), start + 6)):
                combined += normalized_words[end]
                if combined == alias_key:
                    return alias, row[end].right
                if len(combined) > len(alias_key) + 4:
                    break

    return None


def word_looks_like_any_alias(text: str) -> bool:
    normalized = normalize_compare_value(text)
    if not normalized:
        return False
    return any(
        normalized == normalize_compare_value(alias)
        for definition in FIELD_DEFINITIONS
        for alias in definition.aliases
    )


def median_number(values: list[int]) -> float:
    if not values:
        return 0

    ordered = sorted(values)
    middle = len(ordered) // 2
    if len(ordered) % 2:
        return float(ordered[middle])

    return (ordered[middle - 1] + ordered[middle]) / 2


def build_photo_ocr_variants(image, image_enhance, image_filter, image_ops) -> list:
    base = resize_for_ocr(image.convert("RGB"))
    variants = []
    seen_sizes: set[tuple[int, int, int]] = set()

    def add_variant(variant) -> None:
        key = (variant.width, variant.height, len(variants))
        if variant.width < 120 or variant.height < 80:
            return
        if key in seen_sizes:
            return
        seen_sizes.add(key)
        variants.append(variant)

    gray = image_ops.grayscale(base)
    contrast = image_ops.autocontrast(gray)
    enhanced = image_enhance.Contrast(contrast).enhance(1.8)
    sharpened = image_enhance.Sharpness(enhanced).enhance(2.0)
    denoised = sharpened.filter(image_filter.MedianFilter(size=3))
    threshold = estimate_binary_threshold(denoised)
    binary = denoised.point(lambda pixel: 255 if pixel > threshold else 0)

    add_variant(base)
    add_variant(denoised)
    add_variant(binary)

    content_crop = crop_to_dark_content(base)
    if content_crop is not None:
        add_photo_region_variants(content_crop, variants, seen_sizes, image_enhance, image_filter, image_ops)

    for box in locate_photo_mark_boxes(base, image_filter, image_ops)[:4]:
        region = base.crop(box)
        add_photo_region_variants(region, variants, seen_sizes, image_enhance, image_filter, image_ops)

    rotated = [
        base.rotate(angle, expand=True, fillcolor=(255, 255, 255))
        for angle in (-2, 2)
    ]
    for rotated_base in rotated:
        rotated_crop = crop_to_dark_content(rotated_base)
        if rotated_crop is not None:
            add_photo_region_variants(rotated_crop, variants, seen_sizes, image_enhance, image_filter, image_ops)

    return variants[:12]


def add_photo_region_variants(region, variants: list, seen_sizes: set[tuple[int, int, int]], image_enhance, image_filter, image_ops) -> None:
    def add_variant(variant) -> None:
        key = (variant.width, variant.height, len(variants))
        if variant.width < 120 or variant.height < 80:
            return
        if key in seen_sizes:
            return
        seen_sizes.add(key)
        variants.append(variant)

    region = resize_for_ocr(region.convert("RGB"))
    crop_gray = image_ops.grayscale(region)
    crop_contrast = image_ops.autocontrast(crop_gray)
    crop_enhanced = image_enhance.Contrast(crop_contrast).enhance(2.2)
    crop_sharp = image_enhance.Sharpness(crop_enhanced).enhance(2.4)
    crop_denoised = crop_sharp.filter(image_filter.MedianFilter(size=3))
    crop_threshold = estimate_binary_threshold(crop_denoised)
    crop_binary = crop_denoised.point(lambda pixel: 255 if pixel > crop_threshold else 0)

    add_variant(region)
    add_variant(crop_denoised)
    add_variant(crop_binary)


def locate_photo_mark_boxes(image, image_filter, image_ops) -> list[tuple[int, int, int, int]]:
    width, height = image.size
    if width <= 0 or height <= 0:
        return []

    max_analysis_width = 1200
    scale = min(1.0, max_analysis_width / max(width, 1))
    analysis_image = image.resize((int(width * scale), int(height * scale))) if scale < 1 else image
    gray = image_ops.autocontrast(image_ops.grayscale(analysis_image))

    masks = []
    edge_mask = gray.filter(image_filter.FIND_EDGES)
    edge_threshold = estimate_highlight_threshold(edge_mask)
    masks.append(edge_mask.point(lambda pixel: 255 if pixel > edge_threshold else 0))

    dark_threshold = estimate_binary_threshold(gray)
    masks.append(gray.point(lambda pixel: 255 if pixel < dark_threshold else 0))

    candidate_boxes = []
    for mask in masks:
        grouped = mask.filter(image_filter.MaxFilter(31)).filter(image_filter.MinFilter(7))
        candidate_boxes.extend(find_connected_boxes(
            grouped,
            min_width_ratio=0.08,
            min_height_ratio=0.08,
            min_area_ratio=0.008,
            max_area_ratio=0.72,
        ))

    boxes = []
    for box in filter_plausible_photo_mark_boxes(candidate_boxes, analysis_image.width, analysis_image.height):
        boxes.append(expand_box(
            (
                int(box[0] / scale),
                int(box[1] / scale),
                int(box[2] / scale),
                int(box[3] / scale),
            ),
            width,
            height,
            padding=max(30, int(min(width, height) * 0.025)),
        ))

    return dedupe_boxes(sorted(
        boxes,
        key=lambda box: (box[2] - box[0]) * (box[3] - box[1]),
        reverse=True,
    ))


def estimate_highlight_threshold(gray_image) -> int:
    histogram = gray_image.histogram()
    total = sum(histogram)
    if not total:
        return 28

    weighted_sum = sum(index * count for index, count in enumerate(histogram))
    mean = weighted_sum / total
    return max(18, min(55, int(mean * 1.35)))


def filter_plausible_photo_mark_boxes(
    boxes: list[tuple[int, int, int, int]],
    image_width: int,
    image_height: int,
) -> list[tuple[int, int, int, int]]:
    filtered = []

    for box in boxes:
        width = box[2] - box[0]
        height = box[3] - box[1]
        if width <= 0 or height <= 0:
            continue

        area_ratio = (width * height) / max(image_width * image_height, 1)
        aspect_ratio = width / max(height, 1)
        if area_ratio < 0.02 or area_ratio > 0.78:
            continue
        if aspect_ratio < 0.35 or aspect_ratio > 4.8:
            continue
        if width < image_width * 0.12 or height < image_height * 0.10:
            continue

        filtered.append(box)

    return filtered


def resize_for_ocr(image):
    width, height = image.size
    if width <= 0 or height <= 0:
        return image

    target_width = width
    if width < 1800:
        target_width = 1800
    elif width > 2800:
        target_width = 2800

    if target_width == width:
        return image

    target_height = max(1, int(height * target_width / width))
    return image.resize((target_width, target_height))


def estimate_binary_threshold(gray_image) -> int:
    histogram = gray_image.histogram()
    total = sum(histogram)
    if not total:
        return 170

    weighted_sum = sum(index * count for index, count in enumerate(histogram))
    background_mean = weighted_sum / total
    return max(120, min(205, int(background_mean * 0.82)))


def crop_to_dark_content(image):
    width, height = image.size
    if width <= 0 or height <= 0:
        return None

    analysis = image.convert("L")
    threshold = estimate_binary_threshold(analysis)
    mask = analysis.point(lambda pixel: 255 if pixel < threshold else 0)
    box = mask.getbbox()
    if not box:
        return None

    left, top, right, bottom = expand_box(box, width, height, padding=max(24, int(min(width, height) * 0.03)))
    crop_width = right - left
    crop_height = bottom - top
    if crop_width < width * 0.25 or crop_height < height * 0.18:
        return None
    if crop_width > width * 0.96 and crop_height > height * 0.96:
        return None

    return image.crop((left, top, right, bottom))


def merge_ocr_text_outputs(texts: list[str]) -> str:
    seen = set()
    lines = []
    for text in texts:
        for line in normalize_ocr_lines(text):
            key = normalize_compare_value(line)
            if not key or key in seen:
                continue
            seen.add(key)
            lines.append(line)

    return "\n".join(lines)


def extract_fields(text: str, *, source: str) -> list[CartonMarkExtractedField]:
    fields: dict[str, CartonMarkExtractedField] = {}
    normalized_lines = normalize_ocr_lines(text)

    for definition in FIELD_DEFINITIONS:
        value = find_field_value(normalized_lines, definition)
        if value:
            fields[definition.key] = CartonMarkExtractedField(
                key=definition.key,
                label=definition.label,
                value=normalize_extracted_field_value(value, definition),
                confidence=0.76,
                source=source,
            )

    if "wall_type" not in fields:
        wall_match = re.search(r"\b(SINGLE|DOUBLE|TRIPLE)\s+WALL\b", text, re.IGNORECASE)
        if wall_match:
            fields["wall_type"] = CartonMarkExtractedField(
                key="wall_type",
                label=FIELD_BY_KEY["wall_type"].label,
                value=re.sub(r"\s+", " ", wall_match.group(0)).upper(),
                confidence=0.82,
                source=source,
            )

    if "measurement" not in fields:
        measurement_match = re.search(
            r"\b\d+(?:[.,]\d+)?\s*[X×]\s*\d+(?:[.,]\d+)?\s*[X×]\s*\d+(?:[.,]\d+)?\s*(?:CM|CMS|IN|INCHES)?\b",
            text,
            re.IGNORECASE,
        )
        if measurement_match:
            fields["measurement"] = CartonMarkExtractedField(
                key="measurement",
                label=FIELD_BY_KEY["measurement"].label,
                value=re.sub(r"\s+", " ", measurement_match.group(0)).strip(),
                confidence=0.68,
                source=source,
            )

    # Certification stamps are compact tables and OCR often returns their
    # values in reading order while slightly damaging the tiny row labels.
    # A structured stamp parser is more reliable than letting SIZE LIMIT absorb
    # all four adjacent values as one field.
    for key, value in extract_carton_certification_values(text).items():
        definition = FIELD_BY_KEY[key]
        fields[key] = CartonMarkExtractedField(
            key=key,
            label=definition.label,
            value=value,
            confidence=0.78,
            source=source,
        )

    if "barcode" not in fields:
        barcode = find_barcode_candidate(text)
        if barcode:
            fields["barcode"] = CartonMarkExtractedField(
                key="barcode",
                label=FIELD_BY_KEY["barcode"].label,
                value=barcode,
                confidence=0.62,
                source=source,
            )

    return [fields[key] for key in FIELD_ORDER if key in fields]


def extract_carton_certification_values(text: str) -> dict[str, str]:
    normalized = re.sub(r"\s+", " ", str(text or "")).upper()
    markers = (
        "SINGLE WALL", "DOUBLE WALL", "TRIPLE WALL", "SINGLEWALL", "DOUBLEWALL", "TRIPLEWALL",
        "BOX MEET", "CONSTRUCTION",
        "FREIGHT CLASS", "BURSTING", "BUASTING", "MIN.COMB", "MIN COMB",
        "WT.FAC", "WTFAC", "SIZE LIMIT", "GROSS WT", "WT.LT",
    )
    marker_positions = [normalized.find(marker) for marker in markers if normalized.find(marker) >= 0]
    if len(marker_positions) < 2:
        return {}
    stamp_text = normalized[min(marker_positions):]

    values: dict[str, str] = {}
    wall_match = re.search(r"\b(SINGLE|DOUBLE|TRIPLE)\s*WALL\b", stamp_text)
    if wall_match:
        values["wall_type"] = f"{wall_match.group(1)} WALL"

    numeric_values = [int(value) for value in re.findall(r"(?<![A-Z0-9.])([0-9]{2,3})(?![A-Z0-9.])", stamp_text)]
    best_sequence: tuple[int, int, int, int] | None = None
    for index in range(max(0, len(numeric_values) - 3)):
        candidate = tuple(numeric_values[index:index + 4])
        if len(candidate) != 4:
            continue
        bursting, combined, size_limit, gross_limit = candidate
        if 100 <= bursting <= 400 and 20 <= combined <= 150 and 20 <= size_limit <= 120 and 20 <= gross_limit <= 120:
            best_sequence = candidate
            break
    if best_sequence:
        bursting, combined, size_limit, gross_limit = best_sequence
        values.update({
            "bursting_test": str(bursting),
            "min_combined_weight": str(combined),
            "size_limit": str(size_limit),
            "gross_weight_limit": str(gross_limit),
        })
    return values


def enrich_fields_from_expected_values(
    fields: list[CartonMarkExtractedField],
    expected_fields: list[CartonMarkExtractedField],
    text: str,
    *,
    source: str,
) -> list[CartonMarkExtractedField]:
    merged = {field.key: field for field in fields}
    if not normalize_compare_value(text):
        return fields

    for expected in expected_fields:
        if expected.key in merged:
            continue

        expected_value = expected.value.strip()
        if not photo_text_matches_expected_value(text, expected_value, expected.key):
            continue

        definition = FIELD_BY_KEY.get(expected.key)
        if not definition:
            continue

        merged[expected.key] = CartonMarkExtractedField(
            key=definition.key,
            label=definition.label,
            value=expected_value,
            confidence=0.58,
            source=source,
        )

    return [merged[key] for key in FIELD_ORDER if key in merged]


def enrich_expected_fields_from_actual_values(
    expected_fields: list[CartonMarkExtractedField],
    actual_fields: list[CartonMarkExtractedField],
    template_text: str,
    *,
    source: str,
) -> list[CartonMarkExtractedField]:
    merged = {field.key: field for field in expected_fields}
    if not normalize_compare_value(template_text):
        return expected_fields

    for actual in actual_fields:
        if actual.key in merged:
            continue

        actual_value = actual.value.strip()
        if not actual_value:
            continue
        if not photo_text_matches_expected_value(template_text, actual_value, actual.key):
            continue

        definition = FIELD_BY_KEY.get(actual.key)
        if not definition:
            continue

        merged[actual.key] = CartonMarkExtractedField(
            key=definition.key,
            label=definition.label,
            value=actual_value,
            confidence=min(0.62, actual.confidence),
            source=source,
        )

    return [merged[key] for key in FIELD_ORDER if key in merged]


def photo_text_matches_expected_value(text: str, expected_value: str, key: str) -> bool:
    normalized_expected = normalize_compare_value(expected_value)
    if not is_searchable_expected_value(key, normalized_expected):
        return False

    normalized_text = normalize_compare_value(text)
    if normalized_expected in normalized_text:
        return True
    if key not in IDENTIFIER_FIELD_KEYS:
        return False

    numeric_bias = key in {"po", "pi_no", "item", "carton_no", "barcode"}
    expected_identifier = normalize_identifier_value(expected_value, numeric_bias=numeric_bias)
    if not is_searchable_expected_value(key, expected_identifier):
        return False

    candidates = build_identifier_candidates(text, numeric_bias=numeric_bias)
    if any(expected_identifier in candidate for candidate in candidates):
        return True

    max_distance = 2 if len(expected_identifier) >= 8 else 1
    return any(
        fuzzy_contains_identifier(candidate, expected_identifier, max_distance)
        for candidate in candidates
    )


def is_searchable_expected_value(key: str, normalized_value: str) -> bool:
    if key == "barcode":
        return len(normalized_value) >= 8
    if key in {"po", "pi_no", "item", "sku", "carton_no"}:
        return len(normalized_value) >= 4

    return len(normalized_value) >= 6


def normalize_identifier_value(value: str, *, numeric_bias: bool) -> str:
    normalized = value.upper()
    if numeric_bias:
        normalized = normalized.translate(OCR_IDENTIFIER_TRANSLATION)
    return re.sub(r"[^A-Z0-9]+", "", normalized)


def build_identifier_candidates(text: str, *, numeric_bias: bool) -> list[str]:
    tokens = [
        normalize_identifier_value(token, numeric_bias=numeric_bias)
        for token in re.findall(r"[A-Z0-9|]+", text.upper())
    ]
    tokens = [token for token in tokens if token]

    candidates: list[str] = []
    seen: set[str] = set()

    def add_candidate(candidate: str) -> None:
        if candidate and candidate not in seen:
            seen.add(candidate)
            candidates.append(candidate)

    for token in tokens:
        add_candidate(token)

    max_window_size = min(6, len(tokens))
    for window_size in range(2, max_window_size + 1):
        for start in range(0, len(tokens) - window_size + 1):
            add_candidate("".join(tokens[start:start + window_size]))

    joined = "".join(tokens)
    if len(joined) <= 300:
        add_candidate(joined)

    return candidates


def fuzzy_contains_identifier(candidate: str, expected: str, max_distance: int) -> bool:
    if not candidate or not expected:
        return False
    if expected in candidate:
        return True

    min_length = max(1, len(expected) - max_distance)
    max_length = len(expected) + max_distance
    for length in range(min_length, max_length + 1):
        if length > len(candidate):
            continue
        for start in range(0, len(candidate) - length + 1):
            segment = candidate[start:start + length]
            if levenshtein_distance_at_most(expected, segment, max_distance):
                return True

    return False


def levenshtein_distance_at_most(left: str, right: str, max_distance: int) -> bool:
    if abs(len(left) - len(right)) > max_distance:
        return False

    previous = list(range(len(right) + 1))
    for left_index, left_char in enumerate(left, start=1):
        current = [left_index]
        row_min = current[0]
        for right_index, right_char in enumerate(right, start=1):
            cost = 0 if left_char == right_char else 1
            current.append(min(
                current[right_index - 1] + 1,
                previous[right_index] + 1,
                previous[right_index - 1] + cost,
            ))
            row_min = min(row_min, current[-1])

        if row_min > max_distance:
            return False
        previous = current

    return previous[-1] <= max_distance


def levenshtein_distance(left: str, right: str) -> int:
    previous = list(range(len(right) + 1))
    for left_index, left_char in enumerate(left, start=1):
        current = [left_index]
        for right_index, right_char in enumerate(right, start=1):
            cost = 0 if left_char == right_char else 1
            current.append(min(
                current[right_index - 1] + 1,
                previous[right_index] + 1,
                previous[right_index - 1] + cost,
            ))
        previous = current

    return previous[-1]


def merge_metadata_fields(
    fields: list[CartonMarkExtractedField],
    metadata: dict[str, str],
    *,
    source: str,
) -> list[CartonMarkExtractedField]:
    merged = {field.key: field for field in fields}

    for key, value in metadata.items():
        normalized_value = value.strip()
        if not normalized_value or key in merged or key not in FIELD_BY_KEY:
            continue
        definition = FIELD_BY_KEY[key]
        merged[key] = CartonMarkExtractedField(
            key=definition.key,
            label=definition.label,
            value=normalized_value,
            confidence=0.9,
            source=source,
        )

    return [merged[key] for key in FIELD_ORDER if key in merged]


SIDE_WEIGHTED_FIELD_KEYS = {
    "description",
    "quantity",
    "carton_no",
    "box_no",
    "size",
    "gw",
    "nw",
    "measurement",
    "section",
    "pi_no",
    "wall_type",
    "bursting_test",
    "min_combined_weight",
    "size_limit",
    "gross_weight_limit",
}

CERTIFICATION_FIELD_KEYS = {
    "wall_type",
    "bursting_test",
    "min_combined_weight",
    "size_limit",
    "gross_weight_limit",
}

COUNTRY_OF_ORIGIN_NAMES = (
    "UNITED STATES",
    "BANGLADESH",
    "CAMBODIA",
    "INDONESIA",
    "THAILAND",
    "VIETNAM",
    "MALAYSIA",
    "PHILIPPINES",
    "MEXICO",
    "CHINA",
    "INDIA",
    "TAIWAN",
    "USA",
)


def extract_country_of_origin_value(text: str) -> str:
    canonical = normalize_alias_key(str(text or ""))
    for country in COUNTRY_OF_ORIGIN_NAMES:
        country_key = normalize_alias_key(country)
        if f"MADEIN{country_key}" in canonical or f"COUNTRYOFORIGIN{country_key}" in canonical:
            return country
    return ""


def extract_photo_country_panel_evidence(text: str) -> tuple[int, list[str]]:
    panel_matches = [int(value) for value in PHOTO_PANEL_COUNT_RE.findall(str(text or ""))]
    panel_count = max(panel_matches, default=0)
    marked_origins = [
        origin.strip().upper()
        for _index, origin in PHOTO_COUNTRY_PANEL_RE.findall(str(text or ""))
        if origin.strip()
    ]
    if marked_origins:
        return max(panel_count, len(marked_origins)), marked_origins

    origins: list[str] = []
    for line in normalize_ocr_lines(text):
        origin = extract_country_of_origin_value(line)
        if origin and origin not in origins:
            origins.append(origin)
    return max(panel_count, 1 if origins else 0), origins


def compare_country_origin_panels(
    upload_side: str,
    *,
    photo_text: str,
    front_template_text: str,
    side_template_text: str,
) -> list[CartonMarkComparisonItem]:
    expected_faces = [
        ("front", extract_country_of_origin_value(front_template_text)),
        ("side", extract_country_of_origin_value(side_template_text)),
    ]
    expected_faces = [(side, value) for side, value in expected_faces if value]
    if not expected_faces:
        return []

    photo_panel_count, actual_origins = extract_photo_country_panel_evidence(photo_text)
    # Require both template faces only when spatial OCR proves that the single
    # upload contains two repeated product panels.  A close-up of one face must
    # not fail for an unphotographed face, and OCR crop duplicates never increase
    # this count because only primary detector boxes produce the marker.
    required_faces = min(
        len(expected_faces),
        max(1, photo_panel_count),
    )
    expected_origins = [value for _side, value in expected_faces[:required_faces]]
    unmatched_actual = list(actual_origins)
    matching_count = 0
    for expected in expected_origins:
        if expected not in unmatched_actual:
            continue
        unmatched_actual.remove(expected)
        matching_count += 1
    expected_value = f"{' / '.join(expected_origins)}（{required_faces} 面）"
    actual_value = (
        f"{' / '.join(actual_origins)}（{len(actual_origins)} 面）"
        if actual_origins
        else "未识别到产地标识"
    )
    if matching_count >= required_faces:
        status = "pass"
        note = ""
    elif actual_origins and any(origin not in expected_origins for origin in actual_origins):
        status = "mismatch"
        note = "现场照片的产地文字与所选 PDF 页面不一致。"
    else:
        status = "missing_actual"
        note = (
            f"照片空间识别到 {max(1, photo_panel_count)} 个箱唛面板，"
            f"但产地标识仅覆盖 {matching_count}/{required_faces} 面。"
        )
    return [CartonMarkComparisonItem(
        side="side" if required_faces > 1 else upload_side,
        field_key="country_of_origin_panels",
        label="产地标识面板",
        comparison_scope="right_value",
        expected=expected_value,
        actual=actual_value,
        status=status,
        confidence=0.88 if actual_origins else 0.45,
        note=note,
    )]


def compare_observed_photo_fields(
    upload_side: str,
    *,
    photo_text: str,
    photo_status: CartonMarkExtractionStatus,
    photo_fields: list[CartonMarkExtractedField],
    front_template_text: str,
    side_template_text: str,
    front_template_fields: list[CartonMarkExtractedField],
    side_template_fields: list[CartonMarkExtractedField],
) -> list[CartonMarkComparisonItem]:
    """Compare only what the photo actually shows, against either mark face.

    A single photo often contains an unfolded carton with both the front and
    side mark.  Conversely, a true single-face photo should not fail merely
    because fields printed on the other face are outside the frame.
    """
    if not photo_fields:
        if not photo_status.ok:
            seed = front_template_fields if upload_side == "front" else side_template_fields
            return compare_side(upload_side, seed, [], photo_status)
        return [CartonMarkComparisonItem(
            side=upload_side,
            field_key="photo_fields",
            label="照片字段",
            comparison_scope="right_value",
            expected="",
            actual="",
            status="review",
            confidence=0.35,
            note="照片 OCR 有文字，但未能可靠提取核对字段，请人工复核。",
        )]

    front_map = {field.key: field for field in front_template_fields}
    side_map = {field.key: field for field in side_template_fields}
    comparisons: list[CartonMarkComparisonItem] = []
    chosen_sides: dict[str, str] = {}
    for actual_field in photo_fields:
        front_expected = front_map.get(actual_field.key)
        side_expected = side_map.get(actual_field.key)
        candidates = [
            ("front", front_expected),
            ("side", side_expected),
        ]
        matching = [
            (side, expected)
            for side, expected in candidates
            if expected and field_values_match(actual_field.key, expected.value, actual_field.value)
        ]
        if matching:
            chosen_side, expected_field = next(
                ((side, field) for side, field in matching if side == upload_side),
                matching[0],
            )
        else:
            preferred_side = "side" if actual_field.key in SIDE_WEIGHTED_FIELD_KEYS else upload_side
            preferred = side_expected if preferred_side == "side" else front_expected
            fallback = front_expected if preferred_side == "side" else side_expected
            expected_field = preferred or fallback
            chosen_side = preferred_side if preferred else ("front" if front_expected else "side")

        chosen_sides[actual_field.key] = chosen_side
        expected = expected_field.value if expected_field else ""
        if not expected:
            status = "missing_expected"
            note = "所选 PDF 页面正唛和侧唛均未识别到该实拍字段。"
            confidence = min(0.45, actual_field.confidence)
        elif field_values_match(actual_field.key, expected, actual_field.value):
            status = "pass"
            note = ""
            confidence = min(expected_field.confidence, actual_field.confidence)
        else:
            status = "mismatch"
            note = f"实拍字段已按{('正唛' if chosen_side == 'front' else '侧唛')}对齐，与 PDF 模板值不一致。"
            confidence = min(expected_field.confidence, actual_field.confidence)

        comparisons.append(CartonMarkComparisonItem(
            side=chosen_side,
            field_key=actual_field.key,
            label=actual_field.label,
            comparison_scope="right_value",
            expected=expected,
            actual=actual_field.value,
            status=status,
            confidence=confidence,
            note=note,
        ))

    # Check a label only when that label was actually seen in the photo.  This
    # avoids treating the unphotographed face as a column of missing labels.
    actual_labels = extract_field_labels(photo_text)
    front_labels = extract_field_labels(front_template_text)
    side_labels = extract_field_labels(side_template_text)
    for key, actual_label in actual_labels.items():
        if key not in chosen_sides:
            continue
        chosen_side = chosen_sides[key]
        expected_label = (front_labels if chosen_side == "front" else side_labels).get(key, "")
        if not expected_label:
            continue
        if (
            key in CERTIFICATION_FIELD_KEYS
            and any(
                item.field_key == key
                and item.comparison_scope == "right_value"
                and item.expected
                and item.actual
                for item in comparisons
            )
        ):
            # Tiny circular-stamp labels are frequently damaged by OCR.  Once
            # the same structured certification row has been compared by value,
            # a second label mismatch is duplicate noise, not new evidence.
            continue
        comparisons.append(CartonMarkComparisonItem(
            side=chosen_side,
            field_key=f"left_label:{key}",
            label=f"左侧字段名 · {FIELD_BY_KEY[key].label}",
            comparison_scope="left_label",
            expected=expected_label,
            actual=actual_label,
            status="pass" if field_label_values_match(expected_label, actual_label) else "mismatch",
            confidence=0.82,
            note="" if field_label_values_match(expected_label, actual_label) else "PDF 模板左侧字段名与实拍不一致。",
        ))

    comparisons.extend(compare_country_origin_panels(
        upload_side,
        photo_text=photo_text,
        front_template_text=front_template_text,
        side_template_text=side_template_text,
    ))

    observed_keys = {field.key for field in photo_fields}
    strong_product_keys = observed_keys.intersection({"item", "sku", "barcode"})
    meaningful_keys = observed_keys.difference({"customer_name", "po"})
    if (
        comparisons
        and not any(item.status in {"mismatch", "missing_expected"} for item in comparisons)
        and (not strong_product_keys or len(meaningful_keys) < 2)
    ):
        comparisons.append(CartonMarkComparisonItem(
            side=upload_side,
            field_key="photo_coverage",
            label="照片覆盖范围",
            comparison_scope="right_value",
            expected="ITEM/SKU/条码及至少一个箱唛内容字段",
            actual="、".join(FIELD_BY_KEY[key].label for key in FIELD_ORDER if key in meaningful_keys),
            status="review",
            confidence=0.4,
            note="照片识别范围不足以确认对应商品及箱唛主体，请补拍或人工复核。",
        ))

    return comparisons


def cap_comparisons_for_unresolved_pdf_page(
    comparisons: list[CartonMarkComparisonItem],
    page_status: CartonMarkExtractionStatus,
) -> list[CartonMarkComparisonItem]:
    if not page_status.requires_review and not (
        page_status.engine == "pdf-page-selection" and not page_status.ok
    ):
        return comparisons
    if not comparisons:
        return [CartonMarkComparisonItem(
            side="front",
            field_key="pdf_page_selection",
            label="PDF 页面",
            comparison_scope="right_value",
            expected="",
            actual="",
            status="review",
            confidence=0.3,
            note=page_status.message,
        )]
    return [CartonMarkComparisonItem(
        side=item.side,
        field_key=item.field_key,
        label=item.label,
        comparison_scope=item.comparison_scope,
        expected=item.expected,
        actual=item.actual,
        status="review",
        confidence=min(item.confidence, 0.45),
        note=page_status.message,
    ) for item in comparisons]


def compare_side(
    side: str,
    template_fields: list[CartonMarkExtractedField],
    photo_fields: list[CartonMarkExtractedField],
    photo_status: CartonMarkExtractionStatus,
) -> list[CartonMarkComparisonItem]:
    template_map = {field.key: field for field in template_fields}
    photo_map = {field.key: field for field in photo_fields}
    comparisons: list[CartonMarkComparisonItem] = []

    for key in FIELD_ORDER:
        expected_field = template_map.get(key)
        actual_field = photo_map.get(key)
        if not expected_field and not actual_field:
            continue

        label = (expected_field or actual_field).label  # type: ignore[union-attr]
        expected = expected_field.value if expected_field else ""
        actual = actual_field.value if actual_field else ""
        confidence = min(
            expected_field.confidence if expected_field else 0.4,
            actual_field.confidence if actual_field else 0.4,
        )

        if not expected:
            status = "missing_expected"
            note = "PDF 模板未识别到该字段。"
        elif not actual:
            status = "review" if not photo_status.ok else "missing_actual"
            note = photo_status.message if not photo_status.ok else "照片未识别到该字段。"
        elif field_values_match(key, expected, actual):
            status = "pass"
            note = ""
        else:
            status = "mismatch"
            note = "PDF 模板值与实拍识别值不一致。"

        comparisons.append(CartonMarkComparisonItem(
            side=side,
            field_key=key,
            label=label,
            comparison_scope="right_value",
            expected=expected,
            actual=actual,
            status=status,
            confidence=confidence,
            note=note,
        ))

    return comparisons


def compare_label_column(
    side: str,
    template_text: str,
    photo_text: str,
    photo_status: CartonMarkExtractionStatus,
) -> list[CartonMarkComparisonItem]:
    """Compare the printed left-column field names, not only their values."""
    expected_labels = extract_field_labels(template_text)
    actual_labels = extract_field_labels(photo_text)
    comparisons: list[CartonMarkComparisonItem] = []

    for definition in FIELD_DEFINITIONS:
        expected = expected_labels.get(definition.key, "")
        if not expected:
            continue

        actual = actual_labels.get(definition.key, "")
        if not actual:
            status = "review" if not photo_status.ok else "missing_actual"
            note = photo_status.message if not photo_status.ok else "照片左侧字段名未识别，或与 PDF 模板字段名不符。"
        elif field_label_values_match(expected, actual):
            status = "pass"
            note = ""
        else:
            status = "mismatch"
            note = "PDF 模板左侧字段名与实拍不一致。"

        comparisons.append(CartonMarkComparisonItem(
            side=side,
            field_key=f"left_label:{definition.key}",
            label=f"左侧字段名 · {definition.label}",
            comparison_scope="left_label",
            expected=expected,
            actual=actual,
            status=status,
            confidence=0.82 if actual else 0.45,
            note=note,
        ))

    return comparisons


def extract_field_labels(text: str) -> dict[str, str]:
    candidates = extract_label_candidates(text)
    labels: dict[str, str] = {}
    used_candidates: set[str] = set()
    for definition in FIELD_DEFINITIONS:
        available = [
            candidate
            for candidate in candidates
            if normalize_alias_key(candidate) not in used_candidates
        ]
        exact = find_exact_field_label(available, definition)
        if exact:
            labels[definition.key] = exact
            used_candidates.add(normalize_alias_key(exact))
            continue

        approximate = find_approximate_field_label(available, definition)
        if approximate:
            labels[definition.key] = approximate
            used_candidates.add(normalize_alias_key(approximate))

    return labels


def extract_label_candidates(text: str) -> list[str]:
    candidates: list[str] = []
    seen = set()
    for line in normalize_ocr_lines(text):
        first_cell = line.split("|", 1)[0].strip(" :：|#.-")
        if re.match(r"^\d", first_cell):
            first_cell = re.sub(r"^\d+(?:[.,]\d+)?", "", first_cell).strip(" :：|#.-")
        else:
            first_cell = re.split(r"\d", first_cell, maxsplit=1)[0].strip(" :：|#.-")
        if not first_cell or not re.search(r"[A-Z一-龥]", first_cell, re.IGNORECASE):
            continue
        key = normalize_alias_key(first_cell)
        if not key or key in seen:
            continue
        seen.add(key)
        candidates.append(first_cell)
    return candidates


def find_exact_field_label(candidates: list[str], definition: FieldDefinition) -> str:
    for candidate in candidates:
        candidate_key = normalize_alias_key(candidate)
        for alias in sorted(definition.aliases, key=lambda value: len(normalize_alias_key(value)), reverse=True):
            if candidate_key == normalize_alias_key(alias):
                return candidate
    return ""


def find_approximate_field_label(candidates: list[str], definition: FieldDefinition) -> str:
    best_candidate = ""
    best_distance: int | None = None
    for candidate in candidates:
        candidate_key = normalize_alias_key(candidate)
        if len(candidate_key) < 5:
            continue
        for alias in definition.aliases:
            alias_key = normalize_alias_key(alias)
            if len(alias_key) < 6:
                continue
            distance_limit = max(2, int(max(len(candidate_key), len(alias_key)) * 0.42))
            if not levenshtein_distance_at_most(candidate_key, alias_key, distance_limit):
                continue
            distance = levenshtein_distance(candidate_key, alias_key)
            if best_distance is None or distance < best_distance:
                best_candidate = candidate
                best_distance = distance
    return best_candidate


def field_label_values_match(expected: str, actual: str) -> bool:
    def canonical(value: str) -> str:
        normalized = normalize_alias_key(value)
        return re.sub(r"(?:OF|DE)$", "", normalized)

    expected_key = canonical(expected)
    actual_key = canonical(actual)
    if expected_key == actual_key:
        return True

    # Equivalent labels (for example ``ITEM`` and ``Item No.``) are layout
    # variants of the same field, not changed carton-mark content.  Keep the
    # comparison generic by deriving equivalence from FIELD_DEFINITIONS rather
    # than maintaining customer-specific pairs.
    for definition in FIELD_DEFINITIONS:
        aliases = {canonical(alias) for alias in definition.aliases}
        if expected_key in aliases and actual_key in aliases:
            return True
    return False


def field_values_match(key: str, expected: str, actual: str) -> bool:
    if normalize_compare_value(expected) == normalize_compare_value(actual):
        return True
    if key == "description":
        expected_key = normalize_compare_value(expected)
        actual_key = normalize_compare_value(actual)
        return min(len(expected_key), len(actual_key)) >= 8 and (
            expected_key in actual_key or actual_key in expected_key
        )
    if key in {"quantity", "gw", "nw", "measurement", "bursting_test", "min_combined_weight", "size_limit", "gross_weight_limit"}:
        expected_numbers = [Decimal(value.replace(",", ".")) for value in re.findall(r"\d+(?:[.,]\d+)?", expected)]
        actual_numbers = [Decimal(value.replace(",", ".")) for value in re.findall(r"\d+(?:[.,]\d+)?", actual)]
        if expected_numbers and expected_numbers == actual_numbers:
            return True
    if key == "carton_no":
        expected_numbers = re.findall(r"\d+", expected)
        actual_numbers = re.findall(r"\d+", actual)
        if expected_numbers and actual_numbers:
            if expected_numbers == actual_numbers:
                return True
            if len(expected_numbers) == 1 and len(actual_numbers) == 1:
                return expected_numbers[0].lstrip("0") == actual_numbers[0].lstrip("0")
    if key in IDENTIFIER_FIELD_KEYS:
        numeric_bias = key in {"po", "pi_no", "item", "carton_no", "barcode"}
        return normalize_identifier_value(expected, numeric_bias=numeric_bias) == normalize_identifier_value(
            actual,
            numeric_bias=numeric_bias,
        )

    return False


def summarize_comparisons(comparisons: list[CartonMarkComparisonItem]) -> CartonMarkAutoCheckSummary:
    pass_count = sum(1 for item in comparisons if item.status == "pass")
    mismatch_count = sum(1 for item in comparisons if item.status == "mismatch")
    missing_count = sum(1 for item in comparisons if item.status in {"missing_expected", "missing_actual"})
    review_count = sum(1 for item in comparisons if item.status == "review")

    if mismatch_count:
        overall_status = "发现异常"
    elif review_count or missing_count:
        overall_status = "需复核"
    elif pass_count:
        overall_status = "核对通过"
    else:
        overall_status = "未识别"

    return CartonMarkAutoCheckSummary(
        overall_status=overall_status,
        pass_count=pass_count,
        mismatch_count=mismatch_count,
        missing_count=missing_count,
        review_count=review_count,
    )


def find_field_value(lines: list[str], definition: FieldDefinition) -> str:
    for index, line in enumerate(lines):
        for alias in sorted(definition.aliases, key=lambda value: len(normalize_compare_value(value)), reverse=True):
            if (
                definition.key == "quantity"
                and normalize_alias_key(alias) == "PCSCTN"
                and not re.match(
                    rf"^\s*(?:\d+(?:[.,]\d+)?\s*)?{build_flexible_alias_pattern(alias)}\b",
                    line,
                    re.IGNORECASE,
                )
            ):
                continue
            if alias_is_shadowed_by_other_field(line, alias, definition):
                continue
            value = extract_value_after_alias(line, alias)
            if is_plausible_field_value(value, definition):
                return clean_field_value(value)
            if len(normalize_compare_value(alias)) >= 6:
                value = extract_value_before_alias(line, alias)
                if is_plausible_field_value(value, definition):
                    return clean_field_value(value)
            if line_looks_like_alias(line, alias):
                next_value = find_next_line_value(lines, index, definition)
                if next_value:
                    return next_value

    return ""


def alias_is_shadowed_by_other_field(line: str, alias: str, definition: FieldDefinition) -> bool:
    """Do not parse a short field alias inside a longer label from another field."""
    alias_key = normalize_alias_key(alias)
    line_key = normalize_alias_key(line)
    if not alias_key or not line_key:
        return False

    for other_definition in FIELD_DEFINITIONS:
        if other_definition.key == definition.key:
            continue
        for other_alias in other_definition.aliases:
            other_key = normalize_alias_key(other_alias)
            if len(other_key) <= len(alias_key):
                continue
            if alias_key in other_key and other_key in line_key:
                return True

    return False


def normalize_alias_key(value: str) -> str:
    return re.sub(r"[^A-Z0-9一-龥]+", "", value.upper())


def normalize_extracted_field_value(value: str, definition: FieldDefinition) -> str:
    cleaned = clean_field_value(value)
    cleaned = re.sub(r"\s+\*?KEYCODE\b.*$", "", cleaned, flags=re.IGNORECASE).strip()
    cleaned = re.sub(r"\s+MADE\s+IN\s+CHINA\b.*$", "", cleaned, flags=re.IGNORECASE).strip()
    if definition.key == "description":
        return re.split(r"\s+(?:GRS\.?\s*WT|D|DIM)\s*[:：]", cleaned, maxsplit=1, flags=re.IGNORECASE)[0].strip()
    if definition.key in {"item", "sku", "pi_no"}:
        token_match = re.search(r"[A-Z0-9][A-Z0-9./\-]*", cleaned, re.IGNORECASE)
        token = token_match.group(0) if token_match else cleaned
        if definition.key in {"item", "sku"}:
            token = re.sub(r"(?:ITEM(?:NO)?|DESC(?:RIPTION)?|CTN|QTY)$", "", token, flags=re.IGNORECASE)
        return token
    if definition.key == "barcode":
        digit_match = re.search(r"\d{8,18}", re.sub(r"[\s-]+", "", cleaned))
        return digit_match.group(0) if digit_match else cleaned
    if definition.key in {"gw", "nw", "bursting_test", "min_combined_weight", "size_limit", "gross_weight_limit"}:
        numeric_match = re.search(
            r"[-+]?\d+(?:[.,]\d+)?(?:\s*(?:KGS?|LBS?|KG\s*S|INCHES?|M\.?(?:SQ\.?)?(?:FT\.?)?))?",
            cleaned,
            re.IGNORECASE,
        )
        return re.sub(r"\s+", " ", numeric_match.group(0)).strip() if numeric_match else cleaned
    if definition.key == "measurement":
        return re.sub(r"^\(?\s*CM\s*\)?\s*[:：-]*\s*", "", cleaned, flags=re.IGNORECASE)
    if definition.key != "quantity":
        return cleaned

    numeric_values = re.findall(r"\d+(?:[.,]\d+)?", cleaned)
    if numeric_values:
        return numeric_values[0]
    return cleaned


def extract_value_after_alias(line: str, alias: str) -> str:
    pattern = re.compile(
        rf"(?<![A-Z0-9]){build_flexible_alias_pattern(alias)}(?![A-Z0-9])\s*(?:[:：#=.\-])?\s*(.*)",
        re.IGNORECASE,
    )
    match = pattern.search(line)
    if not match:
        return ""

    value = match.group(1)
    value = strip_leading_label_noise(value)
    value = truncate_at_next_field_alias(value, alias)

    return value


def extract_value_before_alias(line: str, alias: str) -> str:
    """Recover table cells exported as `value + label` in vector PDFs."""
    match = re.search(
        rf"^(.*?){build_flexible_alias_pattern(alias)}(?![A-Z0-9])",
        line,
        re.IGNORECASE,
    )
    if not match:
        return ""

    prefix = match.group(1).strip(" \t|:/\\#=_;")
    if not prefix:
        return ""

    # Keep only the last cell/token: a row can contain another label before
    # this one, while the immediately preceding token is the actual value.
    prefix = prefix.split("|")[-1]
    tokens = re.findall(r"[A-Z0-9][A-Z0-9./-]*", prefix, re.IGNORECASE)
    return tokens[-1] if tokens else ""


def find_next_line_value(lines: list[str], index: int, definition: FieldDefinition) -> str:
    for next_line in lines[index + 1:index + 4]:
        candidate = clean_field_value(truncate_at_next_field_alias(strip_leading_label_noise(next_line), ""))
        if is_plausible_field_value(candidate, definition):
            return candidate

    return ""


def line_looks_like_alias(line: str, alias: str) -> bool:
    cleaned_line = normalize_compare_value(line)
    cleaned_alias = normalize_compare_value(alias)
    if not cleaned_line or not cleaned_alias:
        return False
    if cleaned_line == cleaned_alias:
        return True
    return bool(re.fullmatch(rf"{build_flexible_alias_pattern(alias)}(?:N[O0]|NUMBER|NUM)?", line.strip(), re.IGNORECASE))


def build_flexible_alias_pattern(alias: str) -> str:
    parts = []
    for char in alias:
        upper = char.upper()
        if upper.isspace() or upper in {".", "/", "-", "_"}:
            parts.append(r"[\s./_\-]*")
        elif upper == "O":
            parts.append("[O0]")
        elif upper == "I":
            parts.append("[I1L]")
        elif upper == "S":
            parts.append("[S5]")
        elif upper == "B":
            parts.append("[B8]")
        elif upper.isalnum():
            parts.append(re.escape(upper))
        else:
            parts.append(re.escape(char))

    return "".join(parts)


def strip_leading_label_noise(value: str) -> str:
    value = value.strip(" \t|:/\\#=_;")
    value = re.sub(r"^(?:N[O0]|NO\.|NUMBER|NUM|#)\b\s*[:：#=.\-]?\s*", "", value, flags=re.IGNORECASE)
    return value.strip(" \t|:/\\#=_;")


def truncate_at_next_field_alias(value: str, current_alias: str) -> str:
    earliest = len(value)
    current_key = normalize_compare_value(current_alias)

    for definition in FIELD_DEFINITIONS:
        for alias in definition.aliases:
            if current_key and normalize_compare_value(alias) == current_key:
                continue
            split_match = re.search(
                rf"(?<![A-Z0-9]){build_flexible_alias_pattern(alias)}(?![A-Z0-9])",
                value,
                re.IGNORECASE,
            )
            if split_match and split_match.start() < earliest:
                earliest = split_match.start()

    return value[:earliest]


def is_plausible_field_value(value: str, definition: FieldDefinition) -> bool:
    cleaned = clean_field_value(value)
    normalized = normalize_compare_value(cleaned)
    if not normalized:
        return False
    if any(normalized == normalize_compare_value(alias) for alias in definition.aliases):
        return False
    if any(normalized == normalize_compare_value(other_alias) for field in FIELD_DEFINITIONS for other_alias in field.aliases):
        return False

    if definition.key in {
        "po", "pi_no", "item", "sku", "quantity", "carton_no", "box_no", "gw", "nw",
        "measurement", "barcode", "bursting_test", "min_combined_weight", "size_limit",
        "gross_weight_limit",
    } and not re.search(r"\d", cleaned):
        return False
    if definition.key == "measurement" and len(re.findall(r"\d+(?:[.,]\d+)?", cleaned)) < 2:
        return False
    if definition.key in {"gw", "nw", "bursting_test", "min_combined_weight", "size_limit", "gross_weight_limit"}:
        return bool(re.search(r"\d+(?:[.,]\d+)?", cleaned))

    return bool(re.search(r"[A-Z0-9一-龥]", cleaned, re.IGNORECASE))


def find_barcode_candidate(text: str) -> str:
    candidates = re.findall(r"\b\d{12,14}\b", text)
    for match in re.finditer(r"(?<!\d)(?:\d[\s-]+){11,13}\d(?!\d)", text):
        candidates.append(re.sub(r"\D", "", match.group(0)))
    candidates = [candidate for candidate in candidates if barcode_candidate_is_safe(candidate, text)]
    if not candidates:
        return ""

    return max(candidates, key=len)


def barcode_candidate_is_safe(candidate: str, context: str) -> bool:
    digits = re.sub(r"\D", "", candidate)
    if len(digits) not in {12, 13, 14}:
        return False
    for match in re.finditer(re.escape(candidate), context, re.IGNORECASE):
        prefix = context[max(0, match.start() - 28):match.start()].upper()
        if re.search(r"(?:ORDER|PO|P\.O|CONTRACT|DOCUMENT|DOC|PI|ITEM|ARTICLE)\s*(?:NO|NUMBER|#)?\W*$", prefix):
            return False
    return True


def normalize_ocr_lines(text: str) -> list[str]:
    text = text.replace("\r", "\n")
    lines = []
    for raw_line in text.split("\n"):
        line = re.sub(r"\s+", " ", raw_line).strip(" \t|")
        if line:
            lines.append(line)
    return lines


def clean_field_value(value: str) -> str:
    value = re.sub(r"\s+", " ", value).strip(" :：|#.-")
    return value[:120]


def normalize_compare_value(value: str) -> str:
    value = value.upper()
    value = value.replace("O", "0") if re.fullmatch(r"[A-Z0-9\s\-/.]+", value) else value
    return re.sub(r"[^A-Z0-9一-龥]+", "", value)


def fallback_decode_bytes(content: bytes) -> str:
    decoded = content.decode("utf-8", errors="ignore") or content.decode("latin-1", errors="ignore")
    return "".join(char if char in TEXT_CHARS or "\u4e00" <= char <= "\u9fff" else " " for char in decoded)


def clip_text(text: str, limit: int = 4000) -> str:
    text = re.sub(r"\s+", " ", text).strip()
    return text[:limit]
