"""Allowlisted empty carton templates and their native workbook layouts.

Static assets are byte-identical copies of public/templates, bundled below app/
because the API image does not contain the frontend public directory. Contract
tests require both copies to move together. No template operation queries a DB.
"""
from copy import copy, deepcopy
from functools import lru_cache
from io import BytesIO
import json
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.styles import Alignment

BUNDLED_TEMPLATES = Path(__file__).resolve().parents[1] / 'templates' / 'spreadsheet_merge'

# User input is an identifier into this catalogue, never a path or generator name.
_CATALOG = {
    'carton-packaging': {
        'label': '货号包装导入模板', 'profile': 'carton-packaging', 'kind': 'configurations',
        'sheet_name': '导入数据', 'header_row': 1, 'audit_sheet': '说明与示例', 'max_data_rows': 1000,
        'description': '按货号及配置组维护包装资料；校验文本货号、尺寸、用量和重量。仅作静态校验，导入时仍须核对本厂已有资料。',
    },
    'carton-paper-options': {
        'label': '纸品选项导入模板', 'profile': 'generic', 'kind': 'paper-options',
        'sheet_name': '导入数据', 'header_row': 1, 'audit_sheet': '说明与示例', 'max_data_rows': 1000,
        'description': '维护纸品类型、纸质及尺寸选项；保留系统原模板。转换结果需进入基础资料的纸品选项导入预览校验。',
    },
    'carton-locations': {
        'label': '仓位导入模板', 'profile': 'generic', 'kind': 'locations',
        'sheet_name': '导入数据', 'header_row': 1, 'audit_sheet': '说明与示例', 'max_data_rows': 1000,
        'description': '维护仓库及仓位或范围，保留文本前导零；范围展开、重复项和已有仓位由原业务导入预览校验。',
    },
    'carton-history-orders': {
        'label': '历史订单导入模板', 'profile': 'generic', 'file': 'carton-history-order-import-template.xlsx',
        'sheet_name': '历史订单导入', 'header_row': 2, 'audit_sheet': '字段说明', 'max_data_rows': 4999,
        'description': '使用客户、合同、货号及各纸品需求数量的历史宽表；保留标题行及空白附加纸品明细。附加纸品须另行核对填写，示例不导入；订单分组与资料匹配须经历史订单导入预览。',
    },
    'carton-opening-inventory': {
        'label': '期初库存导入模板', 'profile': 'generic', 'file': 'carton-history-inventory-import-template.xlsx',
        'sheet_name': '期初库存导入', 'header_row': 2, 'audit_sheet': '字段说明', 'max_data_rows': 4999,
        'description': '登记期初结余，入库时间逐行填写；不能以原入库总量替代结余。客户、仓库、尺寸单位、价格及实际仓位须在原期初导入工作区设置并核对。',
    },
    'carton-weekly-schedule': {
        'label': '周排期核对导入模板', 'profile': 'generic', 'file': 'carton-weekly-schedule-template.xlsx',
        'sheet_name': '导入数据', 'header_row': 3, 'audit_sheet': '填写说明与示例', 'max_data_rows': 4998,
        'description': '保留两行说明及第3行表头；Reference、客户采购单号与PO.NO各有独立含义。排期数量及现有纸箱订单匹配须进入原周排期导入预览核对。',
    },
}


def _spec(template_id):
    try:
        return _CATALOG[template_id]
    except (KeyError, TypeError) as exc:
        raise HTTPException(404, '模板不存在') from exc


def download_template(template_id):
    spec = _spec(template_id)
    if 'kind' in spec:
        from app.services.carton_master_import import template
        return template(spec['kind'])
    try:
        return (BUNDLED_TEMPLATES / spec['file']).read_bytes()
    except OSError as exc:
        raise HTTPException(503, '系统空模板资源不可用，请联系管理员') from exc


@lru_cache(maxsize=6)
def _metadata(template_id):
    spec = _spec(template_id)
    book = load_workbook(BytesIO(download_template(template_id)), read_only=True, data_only=False, keep_links=False)
    try:
        sheet = book[spec['sheet_name']]
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=spec['header_row'], max_row=spec['header_row']))]
        if not headers or any(not isinstance(value, str) or not value.strip() for value in headers):
            raise HTTPException(503, '系统模板表头无效，请联系管理员')
        return {'id': template_id, 'label': spec['label'], 'headers': headers, 'profile': spec['profile'],
                'description': spec['description'], 'sheet_name': spec['sheet_name'], 'header_row': spec['header_row']}
    finally:
        book.close()


def get_template(template_id):
    return deepcopy(_metadata(template_id))


def list_templates():
    return [get_template(template_id) for template_id in _CATALOG]


def audit_sheet_name(template_id):
    return _spec(template_id)['audit_sheet']


def validate_selection(template_id, headers, profile):
    metadata = get_template(template_id)
    if headers != metadata['headers'] or profile != metadata['profile']:
        raise HTTPException(422, '系统模板列名、顺序或校验配置已改变，请重新选择模板')


def create_workbook(template_id, headers, rows):
    """Fill only the canonical input area; examples and auxiliary inputs stay separate."""
    spec = _spec(template_id)
    validate_selection(template_id, headers, spec['profile'])
    if len(rows) > spec['max_data_rows']:
        raise HTTPException(422, f"当前模板单文件最多 {spec['max_data_rows']} 数据行，请启用分批下载")
    book = load_workbook(BytesIO(download_template(template_id)), keep_links=False)
    sheet = book[spec['sheet_name']]
    start = spec['header_row'] + 1
    styles = [copy(sheet.cell(start, column)._style) for column in range(1, len(headers) + 1)]
    # Published static templates have styled blank rows; clear their input area.
    if sheet.max_row >= start:
        sheet.delete_rows(start, sheet.max_row - start + 1)
    for row_index, values in enumerate(rows, start):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column, value)
            cell._style = copy(styles[column - 1])
            if isinstance(value, str):
                cell.data_type = 's'
                cell.number_format = '@'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    return book


def append_audit(book, template_id, records):
    """Store audit as lossless JSON chunks in the canonical non-input sheet.

Business spreadsheet readers materialize every worksheet before filtering it.
Compact records keep their 5001-row/256-column limits, and each JSON cell avoids
creating a second business header. Joining the consecutive cells restores JSONL.
"""
    spec = _spec(template_id)
    sheet = book[spec['audit_sheet']]
    start = max(22, sheet.max_row + 2)  # Header heuristics inspect only the first 20 rows.
    sheet.cell(start, 1, '转换审计：以下连续片段拼接后为 JSONL；仅供核对，不参与业务导入。')
    pending = ''
    row = start + 1
    for record in records:
        value = json.dumps(record, ensure_ascii=False, separators=(',', ':')) + '\n'
        while value:
            take = min(30000 - len(pending), len(value))
            pending += value[:take]
            value = value[take:]
            if len(pending) == 30000:
                if row > 5001:
                    raise HTTPException(413, '审计原文超出原业务工作簿行数限制，请拆分来源后再下载')
                cell = sheet.cell(row, 1, pending)
                cell.data_type = 's'
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                pending = ''
                row += 1
    if pending:
        if row > 5001:
            raise HTTPException(413, '审计原文超出原业务工作簿行数限制，请拆分来源后再下载')
        sheet.cell(row, 1, pending).data_type = 's'
    sheet.column_dimensions['A'].width = max(sheet.column_dimensions['A'].width or 0, 100)
