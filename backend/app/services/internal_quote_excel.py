from __future__ import annotations

from copy import copy
import json
import re
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image as PillowImage, ImageOps

from app.models.internal_quote import (
    InternalQuote,
    InternalQuoteAttachment,
    InternalQuoteSection,
)
from app.schemas.internal_quote import SECTION_CODE_ORDER


P3_TEMPLATE_VERSION = "internal-quote-p3-v1"
P4_TEMPLATE_VERSION = "internal-quote-p4-v2"
WORKBOOK_LAYOUT_VERSION = "internal-quote-unified-desk-v16"
ENGINEERING_WORKBOOK_TEMPLATE_VERSION = "internal-quote-engineering-template-v1"
ENGINEERING_WORKBOOK_TEMPLATE_PATH = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "internal_quote_export_templates"
    / "工程资料模板.xlsx"
)
TEMPLATE_VERSION = P3_TEMPLATE_VERSION
STRUCTURED_DATA_SCHEMA_VERSION = "internal-quote-structured-data-v1"
STRUCTURED_DATA_CHUNK_SIZE = 30000
SECTION_ORDER = SECTION_CODE_ORDER

NAVY = "17324D"
TEAL = "0F766E"
PALE_TEAL = "DFF4F1"
PALE_BLUE = "E8F1F8"
SLATE = "334155"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TEMPLATE_RED = "FF0000"
TEMPLATE_BLUE = "0000FF"
TEMPLATE_BLACK = "000000"
TEMPLATE_YELLOW = "FFF2CC"
TEMPLATE_ORANGE = "FCE4D6"
TEMPLATE_SKY = "DDEBF7"
TEMPLATE_GREEN = "E2F0D9"
TEMPLATE_GRAY = "E7E6E6"
TEMPLATE_THIN = Side(style="thin", color=TEMPLATE_BLACK)
TEMPLATE_MEDIUM = Side(style="medium", color=TEMPLATE_BLACK)
TEMPLATE_BORDER = Border(
    left=TEMPLATE_THIN,
    right=TEMPLATE_THIN,
    top=TEMPLATE_THIN,
    bottom=TEMPLATE_THIN,
)


def _json_object(value: str) -> dict[str, Any]:
    try:
        parsed = json.loads(value or "{}")
    except (TypeError, ValueError):
        return {}
    return parsed if isinstance(parsed, dict) else {}


def _safe_text(value: object) -> str:
    result = "" if value is None else str(value)
    if result.startswith(("=", "+", "-", "@")):
        return f"'{result}"
    return result


def _is_buzzbee_customer(value: object) -> bool:
    normalized = re.sub(r"[^a-z0-9]+", "", str(value or "").lower())
    return "buzzbee" in normalized


def _json_text(value: object) -> str:
    rendered = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return rendered[:32000]


def _canonical_json_text(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _json_chunks(value: object) -> list[str]:
    rendered = _canonical_json_text(value)
    return [
        rendered[index:index + STRUCTURED_DATA_CHUNK_SIZE]
        for index in range(0, len(rendered), STRUCTURED_DATA_CHUNK_SIZE)
    ] or [""]


def _number(value: object) -> float | str:
    try:
        return float(str(value))
    except (TypeError, ValueError):
        return ""


def _float_value(value: object, default: float = 0.0) -> float:
    parsed = _number(value)
    return parsed if isinstance(parsed, float) else default


def _price_number(value: object) -> float | str:
    parsed = _number(value)
    if isinstance(parsed, float):
        return parsed
    match = re.search(r"[-+]?\d[\d,]*(?:\.\d+)?", str(value or ""))
    return float(match.group(0).replace(",", "")) if match else ""


def _dict_value(value: object) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _list_of_dicts(value: object) -> list[dict[str, Any]]:
    return [item for item in value if isinstance(item, dict)] if isinstance(value, list) else []


def _section_payload(section: InternalQuoteSection | None) -> dict[str, Any]:
    return _json_object(section.payload_json) if section is not None else {}


def _section_calculation(section: InternalQuoteSection | None) -> dict[str, Any]:
    return _json_object(section.calculation_json) if section is not None else {}


def _calculation_lines(section: InternalQuoteSection | None, kind: str | None = None) -> list[dict[str, Any]]:
    rows = _list_of_dicts(_section_calculation(section).get("line_breakdown", []))
    return [row for row in rows if row.get("kind") == kind] if kind else rows


def _rr2_values(summary: dict[str, Any], table: str) -> dict[str, float]:
    result: dict[str, float] = {}
    for row in _list_of_dicts(summary.get(table, [])):
        key = str(row.get("key") or "")
        if key:
            result[key] = _float_value(row.get("value"))
    return result


def _rr2_tax_rows(summary: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {
        str(row.get("key")): row
        for row in _list_of_dicts(summary.get("t4", []))
        if row.get("key")
    }


def _material_label(raw_key: object) -> str:
    parts = [part.strip() for part in str(raw_key or "").split("|") if part.strip()]
    if not parts:
        return ""
    material = parts[0]
    grade = parts[1] if len(parts) > 1 else ""
    label = f"{grade}{material}" if grade else material
    return label if label.endswith("料") else f"{label}料"


def _reference_materials(snapshot: dict[str, Any]) -> list[tuple[str, float | str]]:
    source = snapshot.get("material_prices", {})
    rows: list[tuple[str, float | str]] = []
    if isinstance(source, dict):
        for key, value in source.items():
            if isinstance(value, dict):
                price = value.get("price_hkd_lb", value.get("price_hkd", value.get("price", "")))
                label = value.get("label") or _material_label(key)
            else:
                price = value
                label = _material_label(key)
            rows.append((_safe_text(label), _number(price)))
    elif isinstance(source, list):
        for value in _list_of_dicts(source):
            key = value.get("label") or f"{value.get('material', '')}|{value.get('grade', '')}"
            price = value.get("price_hkd_lb", value.get("price_hkd", value.get("price", "")))
            rows.append((_material_label(key), _number(price)))
    return rows


def _reference_machines(snapshot: dict[str, Any]) -> list[tuple[str, float | str]]:
    source = snapshot.get("machine_prices", [])
    rows: list[tuple[str, float | str]] = []
    if isinstance(source, dict):
        source = [{"machine_code": key, "shift_price_hkd": value} for key, value in source.items()]
    for value in _list_of_dicts(source):
        label = (
            value.get("machine_range")
            or value.get("range")
            or value.get("machine_name")
            or value.get("name")
            or value.get("machine_code")
            or value.get("code")
            or ""
        )
        price = value.get("shift_price_hkd", value.get("price_hkd", value.get("price", "")))
        rows.append((_safe_text(label), _number(price)))
    return rows


def _tax_rate(snapshot: dict[str, Any], key: str, default: float) -> float:
    return _float_value(_dict_value(snapshot.get("tax_rates", {})).get(key), default)


def _template_cell(
    sheet,
    row: int,
    column: int,
    value: object = None,
    *,
    bold: bool = False,
    color: str = TEMPLATE_BLACK,
    fill: str | None = None,
    horizontal: str = "center",
    vertical: str = "center",
    wrap_text: bool = True,
    size: int = 10,
    number_format: str | None = None,
    border: Border | None = TEMPLATE_BORDER,
    font_name: str = "宋体",
) -> None:
    cell = sheet.cell(row, column, value)
    cell.font = Font(name=font_name, size=size, bold=bold, color=color)
    cell.alignment = Alignment(
        horizontal=horizontal,
        vertical=vertical,
        wrap_text=wrap_text,
    )
    cell.border = border or Border()
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if number_format:
        cell.number_format = number_format


def _percent_fraction(value: object) -> float:
    """Normalize persisted percent-points to an Excel percentage fraction."""

    numeric = _float_value(value)
    return numeric / 100 if abs(numeric) > 0.5 else numeric


def _template_style_range(
    sheet,
    min_row: int,
    max_row: int,
    min_column: int,
    max_column: int,
    *,
    fill: str | None = None,
    bold: bool = False,
    color: str = TEMPLATE_BLACK,
) -> None:
    for row in range(min_row, max_row + 1):
        for column in range(min_column, max_column + 1):
            _template_cell(
                sheet,
                row,
                column,
                sheet.cell(row, column).value,
                fill=fill,
                bold=bold,
                color=color,
            )


def _replace_border_sides(
    cell,
    *,
    left: Side | None = None,
    right: Side | None = None,
    top: Side | None = None,
    bottom: Side | None = None,
) -> None:
    current = cell.border
    cell.border = Border(
        left=current.left if left is None else left,
        right=current.right if right is None else right,
        top=current.top if top is None else top,
        bottom=current.bottom if bottom is None else bottom,
    )


def _apply_outline_border(
    sheet,
    min_row: int,
    max_row: int,
    min_column: int,
    max_column: int,
    side: Side,
) -> None:
    for column in range(min_column, max_column + 1):
        _replace_border_sides(sheet.cell(min_row, column), top=side)
        _replace_border_sides(sheet.cell(max_row, column), bottom=side)
    for row in range(min_row, max_row + 1):
        _replace_border_sides(sheet.cell(row, min_column), left=side)
        _replace_border_sides(sheet.cell(row, max_column), right=side)


def _apply_table_borders(
    sheet,
    min_row: int,
    max_row: int,
    min_column: int,
    max_column: int,
    *,
    outline: Side = TEMPLATE_MEDIUM,
) -> None:
    """Give every visible table cell a printable grid plus a clear outline."""

    if min_row > max_row or min_column > max_column:
        return
    for row in range(min_row, max_row + 1):
        for column in range(min_column, max_column + 1):
            sheet.cell(row, column).border = TEMPLATE_BORDER
    _apply_outline_border(
        sheet,
        min_row,
        max_row,
        min_column,
        max_column,
        outline,
    )


def _excel_column_pixels(width: float | None) -> int:
    """Approximate Excel column width in pixels using the OOXML convention."""

    value = 8.43 if width is None else max(float(width), 0.0)
    return int(value * 12) if value < 1 else int(value * 7 + 5)


def _excel_row_pixels(height: float | None) -> int:
    points = 15.0 if height is None else max(float(height), 0.0)
    return int(round(points * 96 / 72))


def _latest_product_image_attachment(
    attachments: list[InternalQuoteAttachment] | None,
) -> InternalQuoteAttachment | None:
    candidates = [
        attachment
        for attachment in attachments or []
        if getattr(attachment, "department", "") == "product-image"
        and str(getattr(attachment, "content_type", "")).lower().startswith("image/")
        and bool(getattr(attachment, "content", b""))
    ]
    if not candidates:
        return None
    return max(
        candidates,
        key=lambda item: (
            str(getattr(item, "uploaded_at", "")),
            str(getattr(item, "id", "")),
        ),
    )


def _add_product_image_to_region(
    sheet,
    attachments: list[InternalQuoteAttachment] | None,
    *,
    min_row: int,
    max_row: int,
    min_column: int,
    max_column: int,
) -> bool:
    """Fit the current product image into the blank block above carton data."""

    attachment = _latest_product_image_attachment(attachments)
    if attachment is None or min_row > max_row or min_column > max_column:
        return False

    image_buffer = BytesIO()
    try:
        with PillowImage.open(BytesIO(attachment.content)) as source:
            normalized = ImageOps.exif_transpose(source)
            if normalized.mode not in {"RGB", "RGBA"}:
                normalized = normalized.convert("RGBA")
            normalized.save(image_buffer, format="PNG")
            source_width, source_height = normalized.size
    except (OSError, ValueError):
        return False
    if source_width <= 0 or source_height <= 0:
        return False

    region_width = sum(
        _excel_column_pixels(
            sheet.column_dimensions[get_column_letter(column)].width
        )
        for column in range(min_column, max_column + 1)
    )
    region_height = sum(
        _excel_row_pixels(sheet.row_dimensions[row].height)
        for row in range(min_row, max_row + 1)
    )
    padding = 8
    available_width = max(region_width - padding * 2, 1)
    available_height = max(region_height - padding * 2, 1)
    scale = min(
        available_width / source_width,
        available_height / source_height,
    )
    target_width = max(1, int(round(source_width * scale)))
    target_height = max(1, int(round(source_height * scale)))
    horizontal_offset = max((region_width - target_width) // 2, 0)
    vertical_offset = max((region_height - target_height) // 2, 0)

    image_buffer.seek(0)
    image = OpenpyxlImage(image_buffer)
    image.width = target_width
    image.height = target_height
    image.anchor = OneCellAnchor(
        _from=AnchorMarker(
            col=min_column - 1,
            colOff=pixels_to_EMU(horizontal_offset),
            row=min_row - 1,
            rowOff=pixels_to_EMU(vertical_offset),
        ),
        ext=XDRPositiveSize2D(
            cx=pixels_to_EMU(target_width),
            cy=pixels_to_EMU(target_height),
        ),
    )
    sheet.add_image(image)
    return True


def _positive_integer(value: object) -> int | None:
    numeric = _float_value(value)
    if numeric <= 0 or not numeric.is_integer():
        return None
    return int(numeric)


def _moq_label(value: object) -> str:
    moq = _positive_integer(value)
    if moq is None:
        return "MOQ"
    if moq % 1000 == 0:
        return f"MOQ{moq // 1000}K"
    return f"MOQ{moq:,}"


def _compact_number_label(value: object) -> str:
    numeric = _number(value)
    if not isinstance(numeric, float):
        return _safe_text(value)
    if numeric.is_integer():
        return str(int(numeric))
    return f"{numeric:.4f}".rstrip("0").rstrip(".")


def _matched_reference_price_cell(
    candidates: list[tuple[str, float, str]],
    identifiers: Iterable[object],
    expected_price: object,
) -> str | None:
    normalized_identifiers = [
        re.sub(r"[^0-9A-Z\u3400-\u9FFF]", "", str(value or "").upper()).removesuffix("料")
        for value in identifiers
        if str(value or "").strip()
    ]
    expected = _float_value(expected_price)
    ranked: list[tuple[int, str]] = []
    for label, price, cell_reference in candidates:
        normalized_label = re.sub(
            r"[^0-9A-Z\u3400-\u9FFF]",
            "",
            label.upper(),
        ).removesuffix("料")
        score = sum(
            2
            for identifier in normalized_identifiers
            if identifier and identifier in normalized_label
        )
        if expected > 0 and abs(price - expected) <= max(0.000001, expected * 0.000001):
            score += 4
        if score:
            ranked.append((score, cell_reference))
    return max(ranked, default=(0, ""))[1] or None


def _molding_loss_weight_formula(line: dict[str, Any]) -> str | None:
    loss_rate = _float_value(line.get("loss_rate_percent"))
    loss_weight = _float_value(
        line.get("loss_weight_g", line.get("net_weight_g"))
    )
    net_weight = _float_value(line.get("net_weight_g"))
    if net_weight <= 0 and loss_weight > 0:
        divisor = 1 + loss_rate / 100
        net_weight = loss_weight / divisor if divisor > 0 else loss_weight
    if net_weight > 0:
        return (
            f"={format(net_weight, '.10g')}"
            f"*(1+{format(loss_rate, '.10g')}/100)"
        )
    return None


def _apply_molding_row_formulas(
    sheet: Worksheet,
    row_index: int,
    line: dict[str, Any],
    material_reference_cells: list[tuple[str, float, str]],
    machine_reference_cells: list[tuple[str, float, str]],
) -> None:
    weight_formula = _molding_loss_weight_formula(line)
    if weight_formula:
        sheet.cell(row_index, 5).value = weight_formula

    material_price_reference = _matched_reference_price_cell(
        material_reference_cells,
        (line.get("grade"), line.get("material"), line.get("material_type")),
        line.get("material_price_hkd_lb"),
    )
    material_price_hkd_lb = _float_value(line.get("material_price_hkd_lb"))
    material_price_hkd_g = _float_value(line.get("material_price_hkd_g"))
    material_cost_hkd = _float_value(line.get("material_cost_hkd"))
    if material_price_reference:
        sheet.cell(row_index, 6).value = f"={material_price_reference}/454"
    elif material_price_hkd_lb > 0:
        sheet.cell(row_index, 6).value = (
            f"={format(material_price_hkd_lb, '.10g')}/454"
        )
    elif material_price_hkd_g > 0:
        sheet.cell(row_index, 6).value = f"={format(material_price_hkd_g, '.10g')}"
    elif material_cost_hkd > 0 and weight_formula:
        sheet.cell(row_index, 6).value = (
            f"={format(material_cost_hkd, '.10g')}/E{row_index}"
        )
    sheet.cell(row_index, 6).number_format = "0.0000_ "

    machine_price_reference = _matched_reference_price_cell(
        machine_reference_cells,
        (line.get("machine_code"), line.get("machine_name")),
        line.get("machine_shift_price_hkd"),
    )
    machine_shift_price_hkd = _float_value(line.get("machine_shift_price_hkd"))
    sets = _float_value(line.get("sets", line.get("cavity")))
    target_output = _float_value(line.get("target_output"))
    molding_cost_hkd = _float_value(line.get("molding_cost_hkd"))
    if machine_price_reference:
        sheet.cell(row_index, 10).value = (
            f"={machine_price_reference}/I{row_index}/H{row_index}"
        )
    elif machine_shift_price_hkd > 0:
        sheet.cell(row_index, 10).value = (
            f"={format(machine_shift_price_hkd, '.10g')}"
            f"/I{row_index}/H{row_index}"
        )
    elif molding_cost_hkd > 0 and sets > 0 and target_output > 0:
        derived_shift_price = molding_cost_hkd * sets * target_output
        sheet.cell(row_index, 10).value = (
            f"={format(derived_shift_price, '.10g')}"
            f"/I{row_index}/H{row_index}"
        )
    elif molding_cost_hkd > 0:
        sheet.cell(row_index, 10).value = f"={format(molding_cost_hkd, '.10g')}"
    sheet.cell(row_index, 10).number_format = "0.0000_ "

    if sheet.cell(row_index, 6).value not in (None, ""):
        sheet.cell(row_index, 11).value = f"=E{row_index}*F{row_index}"
    elif material_cost_hkd > 0:
        sheet.cell(row_index, 11).value = f"={format(material_cost_hkd, '.10g')}"
    sheet.cell(row_index, 11).number_format = "0.000_ "


def _pricing_entry_amount(entry: dict[str, Any]) -> object:
    if str(entry.get("section") or "") != "assembly":
        return _float_value(entry.get("amount_hkd"))
    if str(entry.get("kind") or "") not in {
        "assembly_process",
        "assembly_manual_total",
    }:
        return _float_value(entry.get("amount_hkd"))

    persons = _float_value(entry.get("persons", entry.get("total_persons")))
    teams = _float_value(entry.get("teams"), 1.0)
    production_qty = _float_value(entry.get("production_qty"))
    if persons <= 0 or teams <= 0 or production_qty <= 0:
        return _float_value(entry.get("amount_hkd"))

    formula = f"={_compact_number_label(persons)}*$M$4"
    if abs(teams - 1.0) >= 0.0000001:
        formula += f"*{_compact_number_label(teams)}"
    formula += f"/{_compact_number_label(production_qty)}"
    allocation_factor = _float_value(entry.get("formula_allocation_factor"), 1.0)
    if allocation_factor > 0 and abs(allocation_factor - 1.0) >= 0.0000001:
        formula += f"*{format(allocation_factor, '.10g')}"
    return formula


def _pricing_entry_category(entry: dict[str, Any]) -> str:
    section_code = str(entry.get("section") or "")
    kind = str(entry.get("kind") or "")
    category = str(entry.get("category") or "")
    auxiliary = str(entry.get("auxiliary_category") or "")
    label = _safe_text(entry.get("label"))
    if section_code == "engineering":
        if category == "hardware":
            return "五金"
        if category == "packaging":
            return "包装材料"
        return auxiliary or "其他外购"
    if section_code == "electronic":
        return "电子"
    if section_code == "painting":
        return "油漆" if "paint" in kind or "漆" in label else "喷油工"
    if section_code == "slush":
        return "搪胶"
    if section_code == "hair":
        return "车发"
    if section_code == "sewing":
        return "车衣"
    if section_code == "assembly":
        return "装配工" if not re.search(r"包装|pack", label, re.IGNORECASE) else "包装人工"
    if section_code == "sales":
        if kind == "carton":
            return "纸箱"
        return {
            "blister": "吸塑",
            "color_box_inner_card": "彩盒/内卡",
            "leaflet_manual": "利宝/说明书",
        }.get(category, "包装材料")
    return label or "其他成本"


def _pricing_entry_tax_tag(entry: dict[str, Any]) -> str:
    category = _pricing_entry_category(entry)
    if category == "电镀":
        return "¥1%"
    if category == "搪胶":
        return "¥3%"
    if category == "吸塑":
        return "¥6%"
    if category in {"啤工", "喷油工", "装配工", "包装人工", "车衣"}:
        return ""
    return "¥13%"


def _pricing_entry_summary_category(entry: dict[str, Any]) -> str:
    """Return the stable category used by the original management summary.

    Detached multiplier and JustPlay component rows keep their item label in
    column C, while column B must remain one of the established summary
    categories so the visible ``SUMIF`` formulas can aggregate every segment.
    """

    category = _pricing_entry_category(entry)
    section_code = str(entry.get("section") or "")
    if category == "包装人工":
        return "装配工"
    if category == "包装材料":
        return "其他外购"
    if section_code == "engineering" and category not in {
        "五金",
        "吸塑",
        "彩盒/内卡",
        "电池",
        "利宝/说明书",
        "电镀",
        "其他外购",
    }:
        return "其他外购"
    return category


def _assembly_group_detail_rows(
    section: InternalQuoteSection | None,
    expected_total: object,
) -> list[tuple[str, str, str, object]]:
    if section is None or section.calculation_status != "valid":
        return []
    summaries = _list_of_dicts(
        _section_calculation(section).get("group_summaries", [])
    )
    if not summaries:
        return []

    rows: list[tuple[str, str, str, object]] = []
    calculated_total = 0.0
    for index, summary in enumerate(summaries, start=1):
        category = str(summary.get("category") or "assembly").strip().lower()
        if category not in {"assembly", "packaging"}:
            return []
        amount = _float_value(summary.get("amount_hkd_pcs"))
        if amount < 0:
            return []
        calculated_total += amount

        raw_group = str(summary.get("group") or "").strip()
        prefix = "包装" if category == "packaging" else "组装"
        if not raw_group:
            group_label = f"{prefix}{index}"
        elif raw_group.startswith(prefix) or (
            category == "assembly" and raw_group.startswith("装配")
        ):
            group_label = raw_group
        else:
            group_label = f"{prefix}{raw_group}"
        description = (
            f"{group_label}（{_compact_number_label(summary.get('total_persons'))}人/"
            f"{_compact_number_label(summary.get('standard_work_hours'))}h/"
            f"{_compact_number_label(summary.get('production_qty'))}）"
        )
        persons = _float_value(summary.get("total_persons"))
        teams = _float_value(summary.get("teams"), 1.0)
        production_qty = _float_value(summary.get("production_qty"))
        if persons <= 0 or teams <= 0 or production_qty <= 0:
            return []
        formula = f"={_compact_number_label(persons)}*$M$4"
        if abs(teams - 1.0) >= 0.0000001:
            formula += f"*{_compact_number_label(teams)}"
        formula += f"/{_compact_number_label(production_qty)}"
        rows.append(("", "装配工", description, formula))

    expected = _float_value(expected_total)
    tolerance = max(0.001, 0.0001 * len(rows))
    return rows if abs(calculated_total - expected) <= tolerance else []


def _shipping_markup_tiers(
    shipping: dict[str, Any],
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for source in _list_of_dicts(shipping.get("markup_tiers", [])):
        if source.get("include_in_output", True) is False:
            continue
        moq = _positive_integer(source.get("moq"))
        markup = _float_value(source.get("markup"), _float_value(source.get("markup_x")))
        if moq is None or markup <= 0:
            continue
        rows.append(
            {
                "moq": moq,
                "markup": markup,
                "is_active": bool(source.get("is_active")),
            }
        )
    if not rows:
        rows = [
            {
                "moq": _positive_integer(shipping.get("active_markup_moq")) or 10000,
                "markup": _float_value(shipping.get("markup"), 1.2),
                "is_active": True,
            }
        ]
    rows.sort(key=lambda row: int(row["moq"]))
    if not any(bool(row["is_active"]) for row in rows):
        active_moq = _positive_integer(shipping.get("active_markup_moq"))
        selected = next((row for row in rows if row["moq"] == active_moq), rows[-1])
        selected["is_active"] = True
    return rows


def _testing_fee_values(
    sales_section: InternalQuoteSection | None,
    sales_payload: dict[str, Any],
) -> tuple[float, list[dict[str, object]]]:
    if sales_payload.get("testing_fee_enabled", True) is False:
        return 0.0, []
    totals = _dict_value(_section_calculation(sales_section).get("totals", {}))
    total_usd = _float_value(
        totals.get(
            "testing_fee_total_usd",
            sales_payload.get("testing_fee_total_usd"),
        )
    )
    rows = _list_of_dicts(totals.get("testing_fee_tiers", []))
    if not rows:
        raw_moqs = sales_payload.get("testing_fee_moqs")
        if not isinstance(raw_moqs, list):
            legacy_moq = sales_payload.get("testing_fee_moq")
            raw_moqs = [] if legacy_moq in (None, "") else [legacy_moq]
        rows = [
            {
                "moq": value,
                "unit_price_usd": (
                    total_usd / moq
                    if (moq := _positive_integer(value)) is not None
                    else 0
                ),
            }
            for value in raw_moqs
        ]
    result: list[dict[str, object]] = []
    seen: set[int] = set()
    for source in rows:
        moq = _positive_integer(source.get("moq"))
        if moq is None or moq in seen:
            continue
        seen.add(moq)
        result.append(
            {
                "moq": moq,
                "unit_price_usd": _float_value(
                    source.get("unit_price_usd"),
                    total_usd / moq if total_usd else 0.0,
                ),
            }
        )
    result.sort(key=lambda row: int(row["moq"]))
    return total_usd, result


def _dimensions(value: object) -> tuple[float | str, float | str, float | str]:
    source = _dict_value(value)
    return (
        _number(source.get("length", source.get("length_cm", ""))),
        _number(source.get("width", source.get("width_cm", ""))),
        _number(source.get("height", source.get("height_cm", ""))),
    )


def _dimension_unit(value: object) -> str:
    return "cm" if str(value or "").strip().lower() == "cm" else "inch"


def _inch_value_for_unit(value: object, unit: str) -> float | str:
    parsed = _number(value)
    if not isinstance(parsed, float):
        return ""
    return round(parsed * 2.54, 4) if unit == "cm" else parsed


def _dimensions_for_unit(value: object, unit: str) -> tuple[float | str, float | str, float | str]:
    return tuple(_inch_value_for_unit(item, unit) for item in _dimensions(value))


def _plain_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _positive_number(value: object, fallback: float = 0.0) -> float:
    parsed = _float_value(value, fallback)
    return parsed if parsed > 0 else fallback


def _split_engineering_mold_part_names(value: object) -> list[str]:
    source = _plain_text(value).replace("／", "/")
    if not source:
        return []
    result: list[str] = []
    for raw_part in source.split("/"):
        part = raw_part.strip()
        if not part:
            continue
        if "左右" in part:
            result.extend((part.replace("左右", "左"), part.replace("左右", "右")))
        else:
            result.append(part)
    return result


def _engineering_mold_parts(mold: dict[str, Any]) -> list[dict[str, Any]]:
    saved_parts = _list_of_dicts(mold.get("parts", []))
    if saved_parts:
        return [
            {
                "name": _plain_text(part.get("name") or part.get("item")),
                "color": _plain_text(part.get("color")),
                "process": _plain_text(part.get("process")),
                "process_unit_price_hkd": _positive_number(
                    part.get("process_unit_price_hkd")
                ),
                "unit_net_weight_g": _positive_number(
                    part.get("unit_net_weight_g", part.get("net_weight_g"))
                ),
                "output_count": _positive_number(part.get("output_count"), 1.0),
                "quantity": _positive_number(part.get("quantity"), 1.0),
            }
            for part in saved_parts
        ]
    source_name = mold.get("item") or mold.get("name") or mold.get("chinese_name")
    return [
        {
            "name": name,
            "color": "",
            "process": "",
            "process_unit_price_hkd": 0.0,
            "unit_net_weight_g": 0.0,
            "output_count": 1.0,
            "quantity": 1.0,
        }
        for name in _split_engineering_mold_part_names(source_name)
    ]


def _match_key(value: object) -> str:
    return re.sub(r"\s+", "", _plain_text(value)).casefold()


def _mold_material_label(
    mold: dict[str, Any],
    injection: dict[str, Any] | None,
) -> str:
    source = injection or {}
    material = _plain_text(source.get("material") or mold.get("material"))
    grade = _plain_text(source.get("grade") or mold.get("material_type"))
    if material and grade and _match_key(material) not in _match_key(grade):
        return f"{material}-{grade}"
    return grade or material


def _engineering_company_name(quote: InternalQuote) -> str:
    if _match_key(getattr(quote, "factory_id", "")) == "huaxing":
        return "华兴玩具制品(河源)有限公司"
    return (
        _plain_text(getattr(quote, "workshop_name", ""))
        or _plain_text(getattr(quote, "factory_id", ""))
        or "工程资料"
    )


def _merge_and_style(
    sheet,
    cell_range: str,
    value: object = None,
    *,
    bold: bool = False,
    color: str = TEMPLATE_BLACK,
    fill: str | None = None,
    horizontal: str = "center",
    size: int = 10,
) -> None:
    cells = sheet[cell_range]
    for row in cells:
        for cell in row:
            _template_cell(
                sheet,
                cell.row,
                cell.column,
                None,
                bold=bold,
                color=color,
                fill=fill,
                horizontal=horizontal,
                size=size,
            )
    sheet.merge_cells(cell_range)
    top_left = cells[0][0]
    top_left.value = value


def _engineering_mold_groups(
    engineering_payload: dict[str, Any],
    molding_payload: dict[str, Any],
) -> list[dict[str, Any]]:
    molds = _list_of_dicts(engineering_payload.get("molds", []))
    injections = _list_of_dicts(molding_payload.get("injection_lines", []))
    used_injections: set[int] = set()
    groups: list[dict[str, Any]] = []

    for mold_index, mold in enumerate(molds):
        mold_no_key = _match_key(mold.get("mold_no"))
        item_key = _match_key(
            mold.get("item") or mold.get("name") or mold.get("chinese_name")
        )
        matched: list[tuple[int, dict[str, Any]]] = []
        for injection_index, injection in enumerate(injections):
            if injection_index in used_injections:
                continue
            injection_mold_no = _match_key(injection.get("mold_no"))
            injection_item = _match_key(injection.get("item") or injection.get("name"))
            if (
                (mold_no_key and injection_mold_no == mold_no_key)
                or (
                    not mold_no_key
                    and item_key
                    and injection_item == item_key
                )
            ):
                matched.append((injection_index, injection))
                used_injections.add(injection_index)

        parts = _engineering_mold_parts(mold)
        if not parts:
            parts = [
                {
                    "name": _plain_text(
                        mold.get("item")
                        or mold.get("name")
                        or mold.get("chinese_name")
                    ),
                    "color": "",
                    "process": "",
                    "process_unit_price_hkd": 0.0,
                    "unit_net_weight_g": 0.0,
                    "output_count": 1.0,
                    "quantity": 1.0,
                }
            ]
        groups.append(
            {
                "mold": mold,
                "injections": [row for _, row in matched],
                "parts": parts,
                "fallback_mold_no": f"M{mold_index + 1:02d}",
            }
        )

    unmatched_groups: dict[str, list[dict[str, Any]]] = {}
    unmatched_order: list[str] = []
    for injection_index, injection in enumerate(injections):
        if injection_index in used_injections:
            continue
        group_key = _match_key(injection.get("mold_no")) or f"row-{injection_index}"
        if group_key not in unmatched_groups:
            unmatched_groups[group_key] = []
            unmatched_order.append(group_key)
        unmatched_groups[group_key].append(injection)

    for group_key in unmatched_order:
        group_injections = unmatched_groups[group_key]
        first = group_injections[0]
        parts: list[dict[str, Any]] = []
        for injection in group_injections:
            names = _split_engineering_mold_part_names(
                injection.get("item") or injection.get("name")
            ) or [""]
            for name in names:
                parts.append(
                    {
                        "name": name,
                        "color": _plain_text(injection.get("color")),
                        "process": "",
                        "process_unit_price_hkd": 0.0,
                        "unit_net_weight_g": 0.0,
                        "output_count": _positive_number(
                            injection.get("output_count"), 1.0
                        ),
                        "quantity": _positive_number(
                            injection.get("quantity"), 1.0
                        ),
                    }
                )
        groups.append(
            {
                "mold": {
                    "item": first.get("item", ""),
                    "mold_no": first.get("mold_no", ""),
                    "color": first.get("color", ""),
                    "material": first.get("material", ""),
                    "material_type": first.get("grade", ""),
                    "cavity": first.get("cavity", ""),
                    "quantity": first.get("sets", 1),
                    "net_weight_g": first.get("net_weight_g", 0),
                    "machine_code": first.get("machine_code", ""),
                    "target_output": first.get("target_output", 0),
                    "remark": first.get("remark", ""),
                },
                "injections": group_injections,
                "parts": parts,
                "fallback_mold_no": f"M{len(groups) + 1:02d}",
            }
        )
    return groups


def _matching_injection_for_part(
    part: dict[str, Any],
    injections: list[dict[str, Any]],
    part_index: int,
) -> dict[str, Any]:
    part_key = _match_key(part.get("name"))
    if part_key:
        for injection in injections:
            if _match_key(injection.get("item") or injection.get("name")) == part_key:
                return injection
    if not injections:
        return {}
    if len(injections) == 1:
        return injections[0]
    return injections[min(part_index, len(injections) - 1)]


def _build_mold_schedule_sheet(
    workbook: Workbook,
    quote: InternalQuote,
    engineering_section: InternalQuoteSection | None,
    molding_section: InternalQuoteSection | None,
) -> None:
    sheet = workbook.create_sheet("排摸表")
    engineering_payload = _section_payload(engineering_section)
    molding_payload = _section_payload(molding_section)
    groups = _engineering_mold_groups(engineering_payload, molding_payload)

    _merge_and_style(
        sheet,
        "A1:T1",
        _engineering_company_name(quote),
        bold=True,
        color=TEMPLATE_BLUE,
        size=14,
    )
    _merge_and_style(sheet, "A2:T2", "排  模  表", bold=True, size=16)
    _merge_and_style(
        sheet, "A3:B3", f"客户：{_plain_text(getattr(quote, 'customer', ''))}"
    )
    _merge_and_style(
        sheet, "C3:H3", f"产品编号：{_plain_text(getattr(quote, 'quote_no', ''))}"
    )
    _merge_and_style(
        sheet,
        "I3:L3",
        f"产品名称：{_plain_text(getattr(quote, 'product_name', ''))}",
    )
    _merge_and_style(sheet, "M3:N3", "文件编号：", bold=True)
    _merge_and_style(sheet, "O3:Q3", "")
    _merge_and_style(
        sheet, "R3:S3", f"版本：{_plain_text(getattr(quote, 'version_label', ''))}"
    )
    _template_cell(
        sheet,
        3,
        20,
        f"修订：{getattr(quote, 'header_revision', '')}",
        bold=True,
    )
    _merge_and_style(
        sheet,
        "A4:T4",
        "TO：PMC部；啤塑部；物料部；工程部；装配部；品质部",
        horizontal="left",
    )

    headers = (
        "工模编号",
        "配件名称",
        "颜色",
        "色粉编号",
        "PMS",
        "加工内容",
        "加工总单价",
        "整啤净重(g)",
        "原胶件单净重(g)",
        "出模数",
        "用量",
        "用料名称",
        "整啤套数",
        "整啤模腔数",
        "啤机机型",
        "模具日产量",
        "水口比例",
        "配件图片",
        "模具是否放啤",
        "备注",
    )
    for column, header in enumerate(headers, start=1):
        _template_cell(
            sheet,
            5,
            column,
            header,
            bold=True,
            fill=TEMPLATE_SKY,
            size=9,
        )
    sheet.row_dimensions[5].height = 32

    row_index = 6
    for group in groups:
        mold = _dict_value(group.get("mold"))
        injections = _list_of_dicts(group.get("injections"))
        parts = _list_of_dicts(group.get("parts")) or [{}]
        group_start = row_index
        for part_index, part in enumerate(parts):
            injection = _matching_injection_for_part(part, injections, part_index)
            process_price = _positive_number(part.get("process_unit_price_hkd"))
            part_weight = _positive_number(part.get("unit_net_weight_g"))
            shot_weight = _positive_number(
                injection.get("net_weight_g")
                if injection
                else mold.get("net_weight_g")
            )
            output_count = _positive_number(
                part.get("output_count"),
                _positive_number(injection.get("output_count"), 1.0),
            )
            quantity = _positive_number(
                part.get("quantity"),
                _positive_number(injection.get("quantity"), 1.0),
            )
            values = (
                _plain_text(mold.get("mold_no") or group.get("fallback_mold_no")),
                _plain_text(part.get("name")),
                _plain_text(
                    part.get("color")
                    or injection.get("color")
                    or mold.get("color")
                ),
                "",
                "",
                _plain_text(part.get("process") or mold.get("process")),
                process_price or "",
                shot_weight or "",
                part_weight or "",
                output_count,
                quantity,
                _mold_material_label(mold, injection),
                _positive_number(
                    injection.get("sets"),
                    _positive_number(mold.get("quantity"), 1.0),
                ),
                _plain_text(injection.get("cavity") or mold.get("cavity")),
                _plain_text(
                    injection.get("machine_name")
                    or injection.get("machine_code")
                    or mold.get("machine_code")
                ),
                _positive_number(
                    injection.get("target_output"),
                    _positive_number(mold.get("target_output")),
                )
                or "",
                "",
                "",
                "",
                _plain_text(mold.get("remark") or injection.get("remark")),
            )
            for column, value in enumerate(values, start=1):
                _template_cell(
                    sheet,
                    row_index,
                    column,
                    _safe_text(value) if isinstance(value, str) else value,
                    horizontal="left" if column in {2, 6, 12, 20} else "center",
                    size=9,
                    number_format=(
                        "0.0000"
                        if column in {7, 8, 9, 11}
                        and isinstance(value, (int, float))
                        else None
                    ),
                )
            sheet.row_dimensions[row_index].height = 25
            row_index += 1
        group_end = row_index - 1
        if group_end > group_start:
            for column in (1, 8, 12, 13, 14, 15, 16, 17, 18, 19, 20):
                sheet.merge_cells(
                    start_row=group_start,
                    start_column=column,
                    end_row=group_end,
                    end_column=column,
                )

    if row_index == 6:
        for column in range(1, 21):
            _template_cell(sheet, row_index, column, "", size=9)
        sheet.row_dimensions[row_index].height = 25
        row_index += 1

    widths = (10, 30, 13, 13, 11, 18, 13, 13, 15, 10, 10, 18, 11, 13, 13, 13, 12, 12, 13, 24)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 25
    sheet.row_dimensions[2].height = 30
    sheet.row_dimensions[3].height = 24
    sheet.row_dimensions[4].height = 23
    sheet.freeze_panes = "A6"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.print_title_rows = "1:5"
    sheet.print_area = f"A1:T{row_index - 1}"
    _apply_outline_border(sheet, 1, row_index - 1, 1, 20, TEMPLATE_MEDIUM)


def _format_dimension(value: object) -> str:
    parsed = _number(value)
    if not isinstance(parsed, float):
        return ""
    return f"{parsed:g}"


def _purchase_list_rows(
    engineering_payload: dict[str, Any],
    sales_payload: dict[str, Any],
    electronic_payload: dict[str, Any],
) -> list[tuple[object, ...]]:
    result: list[tuple[object, ...]] = []
    for material in _list_of_dicts(engineering_payload.get("materials", [])):
        category = _plain_text(material.get("category"))
        if category == "packaging":
            continue
        purpose = "工程五金" if category == "hardware" else "工程辅料"
        result.append(
            (
                material.get("item", ""),
                material.get("specification", material.get("spec", "")),
                material.get("material", ""),
                material.get("color", ""),
                _number(material.get("quantity", material.get("qty", ""))),
                material.get("supplier", ""),
                material.get("surface_treatment", ""),
                material.get("purpose") or purpose,
                "",
            )
        )

    for packaging in _list_of_dicts(sales_payload.get("packaging_materials", [])):
        result.append(
            (
                packaging.get("item", ""),
                packaging.get("specification", ""),
                packaging.get("material", ""),
                packaging.get("color", ""),
                _number(packaging.get("quantity", "")),
                packaging.get("supplier", ""),
                packaging.get("surface_treatment", ""),
                "业务包装材料",
                "",
            )
        )

    for carton in _list_of_dicts(sales_payload.get("cartons", [])):
        unit = _dimension_unit(carton.get("size_unit"))
        dimensions = tuple(
            _inch_value_for_unit(carton.get(key), unit)
            for key in ("length_in", "width_in", "height_in")
        )
        size = " × ".join(_format_dimension(value) for value in dimensions)
        qty_per_carton = _positive_number(carton.get("qty_per_carton"), 1.0)
        specification = f"{size} {unit}".strip()
        if qty_per_carton:
            specification = f"{specification}；装量 {qty_per_carton:g}".strip("；")
        result.append(
            (
                carton.get("item") or "纸箱",
                specification,
                carton.get("material", ""),
                carton.get("color", ""),
                round(1 / qty_per_carton, 6) if qty_per_carton else "",
                carton.get("supplier", ""),
                carton.get("surface_treatment", ""),
                "纸箱",
                "",
            )
        )

    quote_mode = _plain_text(electronic_payload.get("quote_mode"))
    if quote_mode == "quick":
        electronic_rows = [
            ("", row)
            for row in _list_of_dicts(electronic_payload.get("quick_quotes", []))
        ]
    else:
        electronic_rows = list(
            _walk_components(
                _list_of_dicts(electronic_payload.get("components", []))
            )
        )
    for parent, component in electronic_rows:
        result.append(
            (
                component.get("item") or component.get("name") or parent,
                component.get("specification", component.get("spec", "")),
                component.get("material", ""),
                component.get("color", ""),
                _number(component.get("quantity", 1)),
                component.get("supplier", ""),
                component.get("surface_treatment", ""),
                "电子零件",
                "",
            )
        )
    return result


def _build_purchase_list_sheet(
    workbook: Workbook,
    quote: InternalQuote,
    engineering_section: InternalQuoteSection | None,
    sales_section: InternalQuoteSection | None,
    electronic_section: InternalQuoteSection | None,
) -> None:
    sheet = workbook.create_sheet("外购清单")
    rows = _purchase_list_rows(
        _section_payload(engineering_section),
        _section_payload(sales_section),
        _section_payload(electronic_section),
    )

    _merge_and_style(
        sheet,
        "A1:J1",
        _engineering_company_name(quote),
        bold=True,
        color=TEMPLATE_BLUE,
        size=14,
    )
    _merge_and_style(sheet, "A2:J2", "外 购 件 清 单", bold=True, size=16)
    _merge_and_style(
        sheet, "A3:B3", f"客户：{_plain_text(getattr(quote, 'customer', ''))}"
    )
    _merge_and_style(
        sheet, "C3:E3", f"产品编号：{_plain_text(getattr(quote, 'quote_no', ''))}"
    )
    _merge_and_style(
        sheet,
        "F3:J3",
        f"产品名称：{_plain_text(getattr(quote, 'product_name', ''))}",
    )
    _merge_and_style(
        sheet,
        "A4:D4",
        "TO：PMC部；物料部；工程部；装配部；品质部",
        horizontal="left",
    )
    _merge_and_style(sheet, "E4:F4", "文件编号：", bold=True)
    _template_cell(sheet, 4, 7, "")
    _template_cell(
        sheet,
        4,
        8,
        f"版本：{_plain_text(getattr(quote, 'version_label', ''))}",
    )
    _template_cell(
        sheet,
        4,
        9,
        f"修订：{getattr(quote, 'header_revision', '')}",
    )
    _template_cell(sheet, 4, 10, "")

    headers = (
        "序号",
        "物料名称",
        "物料规格",
        "材料",
        "颜色",
        "用量",
        "供应商",
        "表面处理",
        "用途",
        "单个重量",
    )
    for column, header in enumerate(headers, start=1):
        _template_cell(
            sheet,
            5,
            column,
            header,
            bold=True,
            fill=TEMPLATE_SKY,
            size=10,
        )
    sheet.row_dimensions[5].height = 28

    row_index = 6
    source_rows = rows or [("", "", "", "", "", "", "", "", "")]
    for sequence, source in enumerate(source_rows, start=1):
        values = (sequence if rows else "", *source)
        for column, value in enumerate(values, start=1):
            _template_cell(
                sheet,
                row_index,
                column,
                _safe_text(value) if isinstance(value, str) else value,
                horizontal="left" if column in {2, 3, 4, 7, 8, 9} else "center",
                size=9,
                number_format=(
                    "0.0000"
                    if column == 6 and isinstance(value, (int, float))
                    else None
                ),
            )
        sheet.row_dimensions[row_index].height = 25
        row_index += 1

    widths = (7, 25, 28, 15, 13, 11, 18, 16, 17, 13)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 25
    sheet.row_dimensions[2].height = 30
    sheet.row_dimensions[3].height = 24
    sheet.row_dimensions[4].height = 23
    sheet.freeze_panes = "A6"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.print_title_rows = "1:5"
    sheet.print_area = f"A1:J{row_index - 1}"
    _apply_outline_border(sheet, 1, row_index - 1, 1, 10, TEMPLATE_MEDIUM)


def _capture_template_row(sheet, row: int, end_column: int) -> tuple[list[Any], float | None]:
    return (
        [copy(sheet.cell(row, column)._style) for column in range(1, end_column + 1)],
        sheet.row_dimensions[row].height,
    )


def _apply_template_row(
    sheet,
    row: int,
    template: tuple[list[Any], float | None],
) -> None:
    styles, height = template
    for column, style in enumerate(styles, start=1):
        sheet.cell(row, column)._style = copy(style)
    sheet.row_dimensions[row].height = height


def _populate_engineering_mold_template(
    sheet,
    quote: InternalQuote,
    engineering_section: InternalQuoteSection | None,
    molding_section: InternalQuoteSection | None,
) -> None:
    engineering_payload = _section_payload(engineering_section)
    molding_payload = _section_payload(molding_section)
    groups = _engineering_mold_groups(engineering_payload, molding_payload)

    single_row_template = _capture_template_row(sheet, 17, 20)
    first_row_template = _capture_template_row(sheet, 12, 20)
    middle_row_template = _capture_template_row(sheet, 13, 20)
    last_row_template = _capture_template_row(sheet, 16, 20)
    two_first_template = _capture_template_row(sheet, 6, 20)
    two_last_template = _capture_template_row(sheet, 7, 20)

    for merged_range in list(sheet.merged_cells.ranges):
        if merged_range.min_row >= 6 and merged_range.max_row <= 28:
            sheet.unmerge_cells(str(merged_range))

    required_rows = max(
        1,
        sum(max(1, len(_list_of_dicts(group.get("parts")))) for group in groups),
    )
    capacity = 23
    if required_rows > capacity:
        sheet.insert_rows(29, required_rows - capacity)
        capacity = required_rows

    for row in range(6, 6 + capacity):
        for column in range(1, 21):
            sheet.cell(row, column).value = None

    sheet["A1"] = _engineering_company_name(quote)
    sheet["A3"] = f"客户名称：{_plain_text(getattr(quote, 'customer', ''))}"
    sheet["C3"] = f"产品编号：{_plain_text(getattr(quote, 'quote_no', ''))}"
    sheet["I3"] = f"产品名称：{_plain_text(getattr(quote, 'product_name', ''))}"
    sheet["M3"] = "文件编号："
    sheet["O3"] = ""
    sheet["R3"] = f"版本：{_plain_text(getattr(quote, 'version_label', ''))}"
    sheet["S3"] = f"修订：{getattr(quote, 'header_revision', '')}"

    row_index = 6
    applied_templates: dict[int, tuple[list[Any], float | None]] = {}
    for group in groups:
        mold = _dict_value(group.get("mold"))
        injections = _list_of_dicts(group.get("injections"))
        parts = _list_of_dicts(group.get("parts")) or [{}]
        group_start = row_index
        group_length = len(parts)
        for part_index, part in enumerate(parts):
            if group_length == 1:
                template = single_row_template
            elif group_length == 2:
                template = two_first_template if part_index == 0 else two_last_template
            elif part_index == 0:
                template = first_row_template
            elif part_index == group_length - 1:
                template = last_row_template
            else:
                template = middle_row_template
            _apply_template_row(sheet, row_index, template)
            applied_templates[row_index] = template

            injection = _matching_injection_for_part(part, injections, part_index)
            process_price = _positive_number(part.get("process_unit_price_hkd"))
            part_weight = _positive_number(part.get("unit_net_weight_g"))
            shot_weight = _positive_number(
                injection.get("net_weight_g")
                if injection
                else mold.get("net_weight_g")
            )
            output_count = _positive_number(
                part.get("output_count"),
                _positive_number(injection.get("output_count"), 1.0),
            )
            quantity = _positive_number(
                part.get("quantity"),
                _positive_number(injection.get("quantity"), 1.0),
            )
            values = (
                _plain_text(mold.get("mold_no") or group.get("fallback_mold_no")),
                _plain_text(part.get("name")),
                _plain_text(
                    part.get("color")
                    or injection.get("color")
                    or mold.get("color")
                ),
                "",
                "",
                _plain_text(part.get("process") or mold.get("process")),
                process_price or "",
                shot_weight or "",
                part_weight or "",
                output_count,
                quantity,
                _mold_material_label(mold, injection),
                _positive_number(
                    injection.get("sets"),
                    _positive_number(mold.get("quantity"), 1.0),
                ),
                _plain_text(injection.get("cavity") or mold.get("cavity")),
                _plain_text(
                    injection.get("machine_name")
                    or injection.get("machine_code")
                    or mold.get("machine_code")
                ),
                _positive_number(
                    injection.get("target_output"),
                    _positive_number(mold.get("target_output")),
                )
                or "",
                "",
                "",
                "",
                _plain_text(mold.get("remark") or injection.get("remark")),
            )
            for column, value in enumerate(values, start=1):
                sheet.cell(row_index, column).value = (
                    _safe_text(value) if isinstance(value, str) else value
                )
            row_index += 1

        group_end = row_index - 1
        if group_end > group_start:
            for column in (1, 3, 4, 5, 8, 12, 13, 14, 15, 16, 17):
                sheet.merge_cells(
                    start_row=group_start,
                    start_column=column,
                    end_row=group_end,
                    end_column=column,
                )

    if not groups:
        _apply_template_row(sheet, 6, single_row_template)
    else:
        # openpyxl reconstructs merged-cell borders and can replace the
        # template font/alignment on continuation cells. Reapply the captured
        # template rows after merging so every visible cell keeps the original
        # workbook's exact formatting.
        for row, template in applied_templates.items():
            _apply_template_row(sheet, row, template)


def _populate_engineering_purchase_template(
    sheet,
    quote: InternalQuote,
    engineering_section: InternalQuoteSection | None,
    sales_section: InternalQuoteSection | None,
    electronic_section: InternalQuoteSection | None,
) -> None:
    rows = _purchase_list_rows(
        _section_payload(engineering_section),
        _section_payload(sales_section),
        _section_payload(electronic_section),
    )
    overflow_template = _capture_template_row(sheet, 27, 10)
    capacity = 22
    required_rows = max(1, len(rows))
    if required_rows > capacity:
        sheet.insert_rows(28, required_rows - capacity)
        for row in range(28, 6 + required_rows):
            _apply_template_row(sheet, row, overflow_template)
        capacity = required_rows

    for row in range(6, 6 + capacity):
        for column in range(1, 11):
            sheet.cell(row, column).value = None

    sheet["A1"] = _engineering_company_name(quote)
    sheet["A3"] = f"客户:    {_plain_text(getattr(quote, 'customer', ''))}"
    sheet["C3"] = f"产品货号：{_plain_text(getattr(quote, 'quote_no', ''))}"
    sheet["E3"] = f"产品名称： {_plain_text(getattr(quote, 'product_name', ''))}"
    sheet["G4"] = (
        "文件编号：             "
        f"版本:{_plain_text(getattr(quote, 'version_label', ''))}   "
        f"修订:{getattr(quote, 'header_revision', '')}"
    )

    for sequence, source in enumerate(rows, start=1):
        values = (sequence, *source)
        for column, value in enumerate(values, start=1):
            sheet.cell(sequence + 5, column).value = (
                _safe_text(value) if isinstance(value, str) else value
            )


def build_internal_quote_engineering_workbook(
    quote: InternalQuote,
    sections: list[InternalQuoteSection],
) -> bytes:
    if not ENGINEERING_WORKBOOK_TEMPLATE_PATH.is_file():
        raise FileNotFoundError(
            f"工程资料模板不存在：{ENGINEERING_WORKBOOK_TEMPLATE_PATH}"
        )
    workbook = load_workbook(
        ENGINEERING_WORKBOOK_TEMPLATE_PATH,
        data_only=False,
        read_only=False,
        keep_links=True,
    )
    try:
        by_code = {section.department: section for section in sections}
        mold_sheet = workbook["排摸表"]
        purchase_sheet = workbook["外购清单"]
        # Template photos belong to the reference product. Engineering part
        # images are explicitly optional, so never leak those stale images
        # into a newly generated quote workbook.
        for sheet in workbook.worksheets:
            sheet._images = []
            sheet._charts = []
        _populate_engineering_mold_template(
            mold_sheet,
            quote,
            by_code.get("engineering"),
            by_code.get("molding"),
        )
        _populate_engineering_purchase_template(
            purchase_sheet,
            quote,
            by_code.get("engineering"),
            by_code.get("sales"),
            by_code.get("electronic"),
        )
        workbook.properties.creator = "Royal Regent Nexus"
        workbook.properties.title = f"{quote.quote_no} 工程资料"
        workbook.active = workbook.sheetnames.index("排摸表")
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
    finally:
        workbook.close()


def _style_title(sheet, title: str, end_column: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    cell = sheet.cell(1, 1, title)
    cell.font = Font(name="Microsoft YaHei", size=17, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30


def _header_row(sheet, row: int, headers: Iterable[str]) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row, column, header)
        cell.font = Font(name="Microsoft YaHei", bold=True, color=SLATE)
        cell.fill = PatternFill("solid", fgColor=PALE_TEAL)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _body_row(
    sheet,
    row: int,
    values: Iterable[object],
    *,
    amount_columns: set[int] | None = None,
    text_columns: set[int] | None = None,
) -> None:
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row, column, value)
        cell.font = Font(name="Microsoft YaHei", color=SLATE)
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if text_columns and column in text_columns:
            cell.data_type = "s"
        if amount_columns and column in amount_columns and isinstance(value, (int, float)):
            cell.number_format = "#,##0.0000"


def _finish_sheet(sheet, widths: tuple[int, ...]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A4"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25


def _line_amount(line: dict[str, Any]) -> float | str:
    for key in (
        "amount_hkd",
        "line_hkd",
        "per_piece_hkd",
        "amount_hkd_pcs",
        "deduction_hkd",
        "total_usd",
    ):
        if key in line and line[key] is not None:
            return _number(line[key])
    return ""


def _build_summary_sheet(
    workbook: Workbook,
    quote: InternalQuote,
    sections: list[InternalQuoteSection],
    reference_snapshot: dict[str, Any],
    rr2_cost_summary: dict[str, Any],
    cost_context: dict[str, Any],
    attachments: list[InternalQuoteAttachment] | None = None,
) -> None:
    """Render the current internal-quote template with dynamic detail anchors."""

    sheet = workbook.active
    sheet.title = "报价明细"
    sheet.sheet_view.showGridLines = True
    by_code = {section.department: section for section in sections}
    sales_section = by_code.get("sales")
    sales_payload = _section_payload(sales_section)
    shipping = _dict_value(rr2_cost_summary.get("shipping_pricing", {}))
    t1 = _rr2_values(rr2_cost_summary, "t1")
    t2 = _rr2_values(rr2_cost_summary, "t2")
    t3 = _rr2_values(rr2_cost_summary, "t3")
    t4 = _rr2_tax_rows(rr2_cost_summary)
    sewing_clothes_tax = t4.get("sewcloth13", {})
    sewing_clothes_rate_percent = sewing_clothes_tax.get("rate_percent")
    sewing_clothes_rate = (
        ""
        if sewing_clothes_rate_percent in (None, "")
        else _percent_fraction(sewing_clothes_rate_percent)
    )
    markup_tiers = _shipping_markup_tiers(shipping)
    active_tier = next((row for row in markup_tiers if row.get("is_active")), markup_tiers[-1])
    testing_fee_total_usd, testing_fee_tiers = _testing_fee_values(
        sales_section,
        sales_payload,
    )
    testing_fee_enabled = sales_payload.get("testing_fee_enabled", True) is not False
    freight_output_enabled = shipping.get("freight_enabled", shipping.get("enabled", True)) is not False
    lifting_output_enabled = shipping.get("lifting_enabled", shipping.get("enabled", True)) is not False

    all_shipping_rows = _list_of_dicts(shipping.get("rows", []))
    if all_shipping_rows and str(all_shipping_rows[0].get("name") or "") == "出厂价":
        route_rows = all_shipping_rows[1:]
    else:
        route_rows = all_shipping_rows
    route_start_column = 5
    route_end_column = route_start_column + len(route_rows) - 1
    side_start_column = max(14, route_end_column + 2)
    side_end_column = side_start_column + 3
    outer_right_column = max(18, side_end_column)

    # A1:R6 — latest frozen materials/machines, tax references, carton factors
    # and exchange values.  The extra R column keeps HKD/USD and selected MOQ
    # visible and auditable without crowding the reference template's C:Q grid.
    materials = _reference_materials(reference_snapshot)[:18]
    machines = _reference_machines(reference_snapshot)[:10]
    material_reference_cells: list[tuple[str, float, str]] = []
    machine_reference_cells: list[tuple[str, float, str]] = []
    for row in range(1, 7):
        for column in range(3, outer_right_column + 1):
            _template_cell(
                sheet,
                row,
                column,
                "",
                color=TEMPLATE_RED if row % 2 == 1 else TEMPLATE_BLUE,
                size=10,
                wrap_text=False,
                border=None,
            )
        sheet.row_dimensions[row].height = 12.75 if row < 6 else 13.5

    for header_row, start_index, end_index, start_column in (
        (1, 0, 10, 4),
        (3, 10, 18, 4),
        (5, 0, 10, 4),
    ):
        source = machines if header_row == 5 else materials
        rows = source[start_index:end_index]
        sheet.cell(header_row, 3).value = "机型" if header_row == 5 else "料型"
        sheet.cell(header_row + 1, 3).value = "单价/元" if header_row == 5 else "单价/P"
        for offset, (label, price) in enumerate(rows, start=start_column):
            sheet.cell(header_row, offset).value = label
            sheet.cell(header_row + 1, offset).value = price
            sheet.cell(header_row + 1, offset).number_format = (
                "#,##0" if header_row == 5 else "0.0"
            )
            numeric_price = _float_value(price)
            reference = (
                _safe_text(label),
                numeric_price,
                f"{get_column_letter(offset)}${header_row + 1}",
            )
            if header_row == 5:
                machine_reference_cells.append(reference)
            else:
                material_reference_cells.append(reference)

    fx = _dict_value(reference_snapshot.get("fx", {}))
    sheet["L3"] = "汇率"
    sheet["L4"] = _float_value(fx.get("rmb_hkd"), 0.85)
    sheet["L4"].number_format = "0.00"
    sheet["M3"] = "人工"
    assembly_payload = _section_payload(by_code.get("assembly"))
    assembly_summaries = _list_of_dicts(
        _section_calculation(by_code.get("assembly")).get("group_summaries", [])
    )
    current_labor_base = (
        assembly_payload.get("labor_base_hkd")
        if assembly_payload.get("labor_base_hkd") not in (None, "")
        else assembly_summaries[0].get("labor_base_hkd")
        if assembly_summaries
        else None
    )
    sheet["M4"] = _float_value(
        current_labor_base,
        _float_value(reference_snapshot.get("assembly_labor_base_hkd"), 260),
    )
    sheet["M4"].number_format = "#,##0"
    top_tax_rows = (
        (
            1,
            (
                ("含税1%", _tax_rate(reference_snapshot, "tax_1_percent", 0.0099)),
                ("搪胶类3%", _tax_rate(reference_snapshot, "slush", 0.03)),
                ("车发类5%", _tax_rate(reference_snapshot, "sewing_hair", 0.05)),
                (
                    _safe_text(sewing_clothes_tax.get("label")) or "车衣类",
                    sewing_clothes_rate,
                ),
            ),
        ),
        (
            3,
            (
                ("吸塑类4%", _tax_rate(reference_snapshot, "blister", 0.04)),
                ("运费含税9%", _tax_rate(reference_snapshot, "freight_tax_9", 0.0826)),
                ("含税13%类", _tax_rate(reference_snapshot, "tax_13_percent", 0.115)),
                ("纸箱类", ""),
            ),
        ),
    )
    for header_row, rows in top_tax_rows:
        for column, (label, value) in enumerate(rows, start=14):
            sheet.cell(header_row, column).value = label
            sheet.cell(header_row + 1, column).value = value
            sheet.cell(header_row + 1, column).number_format = "0.00%"

    paper_factor = _float_value(
        sales_payload.get("paper_price_factor"),
        _float_value(reference_snapshot.get("paper_price_factor"), 2.7),
    )
    inner_paper_factor = _float_value(
        sales_payload.get("inner_paper_price_factor"),
        paper_factor,
    )
    flat_card_factor = _float_value(
        sales_payload.get("flat_card_price_factor"),
        paper_factor,
    )
    for column, (label, value, number_format) in enumerate(
        (
            ("主纸箱系数", paper_factor, "0.00"),
            ("内纸箱系数", inner_paper_factor, "0.00"),
            ("平卡系数", flat_card_factor, "0.00"),
            ("杂项", _float_value(shipping.get("misc_ratio"), 0.02), "0.00%"),
        ),
        start=14,
    ):
        sheet.cell(5, column).value = label
        sheet.cell(6, column).value = value
        sheet.cell(6, column).number_format = number_format
    sheet["R3"] = "港币兑美元"
    sheet["R4"] = _float_value(shipping.get("hkd_usd"), _float_value(fx.get("hkd_usd"), 7.8))
    sheet["R4"].number_format = "0.00"
    sheet["R5"] = "本单MOQ"
    sheet["R6"] = int(active_tier["moq"])
    sheet["R6"].number_format = "#,##0"

    # A7:R7 — product quotation title.
    sheet.merge_cells(
        start_row=7,
        start_column=1,
        end_row=7,
        end_column=outer_right_column,
    )
    title = sheet.cell(7, 1)
    title.value = f"{quote.product_name or quote.quote_no}报价"
    title.font = Font(name="宋体", size=14, bold=False, color=TEMPLATE_BLUE)
    title.alignment = Alignment(horizontal="center", vertical="center")
    title.border = Border(
        left=TEMPLATE_MEDIUM,
        right=TEMPLATE_MEDIUM,
        top=TEMPLATE_MEDIUM,
        bottom=TEMPLATE_MEDIUM,
    )
    sheet.row_dimensions[7].height = 27

    # A8:L — molding/injection detail.  At least six rows are kept so the
    # sparse-product sample remains easy to inspect and can be adjusted later.
    mold_headers = (
        "",
        "",
        "名称",
        "料型",
        "料重(G)",
        "料价(G)",
        "机型",
        "1出几套",
        "目标数",
        "啤工",
        "料金额",
        "报价啤工",
    )
    for column, header in enumerate(mold_headers, start=1):
        _template_cell(
            sheet,
            8,
            column,
            header,
            color=TEMPLATE_RED if column not in {1, 2} else TEMPLATE_BLACK,
            wrap_text=False,
        )
    mold_lines = _calculation_lines(by_code.get("molding"), "injection")
    if not mold_lines:
        mold_lines = _list_of_dicts(
            _section_payload(by_code.get("engineering")).get("molds", [])
        )
    mold_slots = max(6, len(mold_lines))
    mold_start_row = 9
    mold_end_row = mold_start_row + mold_slots - 1
    for index in range(mold_slots):
        row_index = mold_start_row + index
        line = mold_lines[index] if index < len(mold_lines) else {}
        values = (
            index + 1 if line else "",
            "",
            _safe_text(line.get("item") or line.get("chinese_name") or line.get("mold_no") or ""),
            _safe_text(line.get("material") or line.get("material_type") or ""),
            _number(line.get("loss_weight_g", line.get("net_weight_g", ""))),
            _number(line.get("material_price_hkd_g", "")),
            _safe_text(line.get("machine_code") or line.get("machine_name") or ""),
            _number(line.get("sets", line.get("cavity", ""))),
            _number(line.get("target_output", "")),
            _number(line.get("molding_cost_hkd", "")),
            _number(line.get("material_cost_hkd", "")),
            _number(line.get("customer_price_hkd", "")),
        )
        for column, value in enumerate(values, start=1):
            _template_cell(
                sheet,
                row_index,
                column,
                value,
                horizontal="left" if column in {2, 3, 4} else "center",
                wrap_text=False,
                number_format=(
                    "0"
                    if column == 1 and value != ""
                    else "0.0000_ "
                    if column in {6, 10, 12}
                    else "0.000_ "
                    if column == 11
                    else None
                ),
                font_name="Times New Roman" if column in {1, 5, 6, 8, 9, 10, 11, 12} else "宋体",
            )
        if line:
            _apply_molding_row_formulas(
                sheet,
                row_index,
                line,
                material_reference_cells,
                machine_reference_cells,
            )
        if line and values[11] == "":
            sheet.cell(row_index, 12).value = f"=J{row_index}*1.15"
            sheet.cell(row_index, 12).number_format = "0.0000_ "
        sheet.row_dimensions[row_index].height = 15
    mold_total_row = mold_end_row + 1
    for column in range(1, 13):
        _template_cell(sheet, mold_total_row, column, "", color=TEMPLATE_RED, wrap_text=False)
    sheet.cell(mold_total_row, 3).value = "模具合计"
    sheet.cell(mold_total_row, 3).font = Font(
        name="宋体",
        size=10,
        bold=True,
        color=TEMPLATE_BLUE,
    )
    for column in (5, 10, 11, 12):
        letter = get_column_letter(column)
        sheet.cell(mold_total_row, column).value = (
            f"=SUM({letter}{mold_start_row}:{letter}{mold_end_row})"
        )
        sheet.cell(mold_total_row, column).number_format = (
            "0.0000_ " if column in {10, 12} else "0.000_ "
        )
    _apply_outline_border(sheet, 8, mold_total_row, 1, 12, TEMPLATE_MEDIUM)

    # A:D — sorted authoritative cost detail.
    detail_rows: list[tuple[str, str, str, object]] = []

    def add_detail(tax_tag: str, category: str, description: str, amount: object) -> None:
        numeric = _float_value(amount)
        if abs(numeric) >= 0.0005:
            detail_rows.append((tax_tag, category, description, numeric))

    def add_formula_detail(
        tax_tag: str,
        category: str,
        description: str,
        formula: str,
    ) -> None:
        detail_rows.append((tax_tag, category, description, formula))

    # The user-facing summary intentionally exposes one molding-material total
    # ("料价").  The tax summaries below are driven by the tax tags on these
    # authoritative detail rows instead of rebuilding categories a second time.
    molding_material_breakdown = _dict_value(
        rr2_cost_summary.get("molding_material_breakdown", {})
    )
    legacy_molding_material = (
        _float_value(t1.get("imp_mat")) + _float_value(t1.get("dom_mat"))
    )
    molding_material = _float_value(
        molding_material_breakdown.get("total_hkd"),
        _float_value(t1.get("material"), legacy_molding_material),
    )
    molding_labor = t3.get("injection_labor", 0.0)
    molding_blow = t1.get("blow", 0.0)
    molding_adjustment = (
        _float_value(cost_context.get("molding_hkd"))
        - molding_material
        - molding_labor
        - molding_blow
    )
    painting_labor = t3.get("painting_labor", 0.0)
    paint_material = t3.get("paint_material", 0.0)
    painting_adjustment = (
        _float_value(cost_context.get("painting_hkd"))
        - painting_labor
        - paint_material
    )
    sewing_hair = t1.get("sewing_hair", 0.0)
    sewing_cloth = t1.get("sewing_cloth", 0.0)
    sewing_cloth_material = min(
        max(_float_value(sewing_clothes_tax.get("amount_hkd")), 0.0),
        max(sewing_cloth, 0.0),
    )
    sewing_cloth_labor = max(sewing_cloth - sewing_cloth_material, 0.0)
    sewing_cloth_material_tax_tag = (
        "¥13%" if sewing_clothes_rate_percent not in (None, "") else ""
    )
    sewing_adjustment = (
        _float_value(cost_context.get("sewing_hkd"))
        - sewing_hair
        - sewing_cloth
    )
    color_box_amount = max(_float_value(t2.get("color_box")), 0.0)

    if mold_lines:
        add_formula_detail("¥13%", "料价", "料价", f"=K{mold_total_row}")
        add_formula_detail("", "啤工", "啤工", f"=J{mold_total_row}")
    else:
        add_detail("¥13%", "料价", "料价", molding_material)
        add_detail("", "啤工", "啤工", molding_labor)
    add_detail("", "啤工", "注塑未分类成本", molding_adjustment)
    assembly_total = (
        _float_value(cost_context.get("assembly_hkd"))
        + _float_value(cost_context.get("packing_labor_hkd"))
    )
    assembly_group_rows = _assembly_group_detail_rows(
        by_code.get("assembly"),
        assembly_total,
    )
    if assembly_group_rows:
        detail_rows.extend(assembly_group_rows)
    else:
        add_detail("", "装配工", "装工", cost_context.get("assembly_hkd"))
        add_detail("", "装配工", "包装", cost_context.get("packing_labor_hkd"))
    add_detail("", "喷油工", "喷油人工", painting_labor)
    add_detail("", "喷油工", "喷油未分类成本", painting_adjustment)
    add_detail("¥13%", "油漆", "油漆", paint_material)
    add_detail("¥3%", "搪胶", "搪胶", cost_context.get("slush_hkd"))
    add_detail("", "吹气", "吹气", molding_blow)
    # Keep every summary category as an explicit detail-table category.  The
    # released workbook uses SUMIF against column B, so classifying sewing and
    # hair as generic purchases both hides their business meaning and breaks
    # the summary formulas.
    add_detail("¥13%", "五金", "五金", t1.get("hardware"))
    add_detail("¥13%", "电子", "电子", t1.get("electronic"))
    add_detail("¥13%", "马达", "马达", t1.get("motor"))
    add_detail("¥6%", "吸塑", "吸塑", t1.get("suction"))
    add_detail("¥13%", "车发", "车发", sewing_hair)
    add_detail(sewing_cloth_material_tax_tag, "车衣", "车衣物料", sewing_cloth_material)
    add_detail("", "车衣", "车衣人工", sewing_cloth_labor)
    add_detail("", "车衣", "车缝未分类成本", sewing_adjustment)
    add_detail("¥13%", "电池", "电池", t2.get("battery"))
    add_detail("¥13%", "利宝/说明书", "利宝/说明书", t2.get("libao"))
    add_detail("¥1%", "电镀", "电镀", t2.get("plating"))
    add_detail("¥13%", "其他外购", "其他外购", t2.get("other_buy"))
    add_detail("¥13%", "其他外购", "胶袋", t1.get("glue_bag"))
    add_detail("¥13%", "彩盒/内卡", "彩盒/内卡", color_box_amount)
    add_detail("", "纸箱", "纸箱", t2.get("carton", cost_context.get("carton_hkd")))

    detail_start_row = mold_total_row + 4
    detail_slots = max(1, len(detail_rows))
    detail_data_end_row = detail_start_row + detail_slots - 1
    for index in range(detail_slots):
        row_index = detail_start_row + index
        values = detail_rows[index] if index < len(detail_rows) else ("", "", "", "")
        for column, value in enumerate(values, start=1):
            _template_cell(
                sheet,
                row_index,
                column,
                value,
                horizontal="left" if column in {2, 3} else "center",
                wrap_text=False,
                size=11 if column in {2, 4} else 10,
                number_format="0.000" if column == 4 else None,
                border=Border(left=TEMPLATE_MEDIUM) if column == 1 else None,
                font_name="Times New Roman" if column == 1 else "宋体",
            )
        sheet.row_dimensions[row_index].height = 15

    # Adaptive right-side packaging/function/color-box/test blocks.  Short
    # details move these blocks upward; full details retain the reference's
    # seven-row visual offset.
    packaging_offset = min(7, max(0, len(detail_rows) - 8))
    packaging_start_row = detail_start_row + packaging_offset
    cartons = _list_of_dicts(sales_payload.get("cartons", []))
    carton = cartons[0] if cartons else {}
    color_box_unit = _dimension_unit(sales_payload.get("color_box_size_unit"))
    dimension_columns = [side_start_column + offset for offset in range(1, 4)]
    dimension_letters = [get_column_letter(column) for column in dimension_columns]
    packaging_rows: list[list[object]] = []
    dimension_row_offsets: set[int] = set()
    carton_row_metadata: list[tuple[int, str]] = []
    flat_card_row_offsets: list[int] = []

    for index, current_carton in enumerate(cartons or [{}]):
        carton_unit = _dimension_unit(current_carton.get("size_unit"))
        carton_dimensions = tuple(
            _inch_value_for_unit(current_carton.get(field, ""), carton_unit)
            for field in ("length_in", "width_in", "height_in")
        )
        row_offset = len(packaging_rows)
        dimension_row_offsets.add(row_offset)
        carton_row_metadata.append((row_offset, carton_unit))
        carton_label = "外箱" if index == 0 else "内箱" if index == 1 else f"内箱{index}"
        packaging_rows.append([f"{carton_label} ({carton_unit}):", *carton_dimensions])

    # Flat-card source dimensions are canonical inches in the quotation
    # payload.  Display them as centimetres here so the exported formula stays
    # auditable in the same shape as the business workbook (cm ÷ 2.54).
    for current_carton in cartons:
        for flat_index, flat_card in enumerate(
            _list_of_dicts(current_carton.get("flat_cards", [])),
            start=1,
        ):
            row_offset = len(packaging_rows)
            dimension_row_offsets.add(row_offset)
            flat_card_row_offsets.append(row_offset)
            flat_name = _safe_text(flat_card.get("name")) or f"平卡{flat_index}"
            packaging_rows.append(
                [
                    f"{flat_name} (cm):",
                    _float_value(flat_card.get("length_in")) * 2.54,
                    _float_value(flat_card.get("width_in")) * 2.54,
                    _number(flat_card.get("quantity", 1)),
                ]
            )

    color_box_offset = len(packaging_rows)
    dimension_row_offsets.add(color_box_offset)
    packaging_rows.append(
        [
            f"彩盒尺寸 ({color_box_unit})",
            *_dimensions_for_unit(
                sales_payload.get("color_box_size_in")
                or sales_payload.get("color_box_size_cm"),
                color_box_unit,
            ),
        ]
    )
    product_offset = len(packaging_rows)
    dimension_row_offsets.add(product_offset)
    packaging_rows.append(
        [
            "产品尺寸 (in)",
            *_dimensions(
                sales_payload.get("product_size_in")
                or sales_payload.get("product_size_cm")
            ),
        ]
    )

    cuft_offset = len(packaging_rows)
    packaging_rows.append(["CUFT:", "", "", ""])
    paperboard_offset = len(packaging_rows)
    packaging_rows.append(["纸板价", "", "", ""])
    carton_price_offset = len(packaging_rows)
    packaging_rows.append(["箱价：", "", "", ""])
    packing_qty_offset = len(packaging_rows)
    packaging_rows.append(["装箱：", _number(carton.get("qty_per_carton", "")), "PCS/1CTN", ""])
    total_offset = len(packaging_rows)
    packaging_rows.append(["合计", "", "", ""])

    def dimension_inputs(row_offset: int, unit: str) -> list[str]:
        absolute_row = packaging_start_row + row_offset
        return [
            f"{dimension_letters[index]}{absolute_row}"
            + ("/2.54" if unit == "cm" else "")
            for index in range(3)
        ]

    outer_inputs = dimension_inputs(*carton_row_metadata[0])
    packaging_rows[cuft_offset][1] = (
        f"={outer_inputs[0]}*{outer_inputs[1]}*{outer_inputs[2]}/1728"
    )

    flat_card_formulas = [
        f"{dimension_letters[0]}{packaging_start_row + row_offset}/2.54"
        f"*{dimension_letters[1]}{packaging_start_row + row_offset}/2.54"
        f"*{dimension_letters[2]}{packaging_start_row + row_offset}*$P$6/1000"
        for row_offset in flat_card_row_offsets
    ]
    packaging_rows[paperboard_offset][1] = (
        "=" + "+".join(flat_card_formulas) if flat_card_formulas else "=0"
    )

    def carton_price_formula(row_offset: int, unit: str, factor_cell: str) -> str:
        inputs = dimension_inputs(row_offset, unit)
        return (
            f"({inputs[0]}+{inputs[1]}+2)"
            f"*({inputs[1]}+{inputs[2]}+1)*{factor_cell}*2/1000"
        )

    outer_carton_formula = carton_price_formula(*carton_row_metadata[0], "$N$6")
    packing_qty_reference = (
        f"{dimension_letters[0]}{packaging_start_row + packing_qty_offset}"
    )
    inner_carton_formulas = [
        f"({carton_price_formula(row_offset, unit, '$O$6')})*{packing_qty_reference}"
        for row_offset, unit in carton_row_metadata[1:]
    ]
    packaging_rows[carton_price_offset][1] = (
        f"={outer_carton_formula}"
        + ("+" + "+".join(inner_carton_formulas) if inner_carton_formulas else "")
    )
    carton_price_reference = (
        f"{dimension_letters[0]}{packaging_start_row + carton_price_offset}"
    )
    packaging_rows[total_offset][1] = (
        f"=IFERROR({carton_price_reference}/{packing_qty_reference},0)"
    )

    for offset, values in enumerate(packaging_rows):
        row_index = packaging_start_row + offset
        for column, value in enumerate(values, start=side_start_column):
            _template_cell(
                sheet,
                row_index,
                column,
                value,
                color=TEMPLATE_BLUE if column == side_start_column and offset == total_offset else TEMPLATE_BLACK,
                horizontal=(
                    "left"
                    if column in {side_start_column, side_start_column + 1}
                    and offset not in dimension_row_offsets
                    else "center"
                ),
                wrap_text=False,
                size=11 if column in dimension_columns and offset in dimension_row_offsets else 10,
                number_format=(
                    "0.000_ "
                    if column == side_start_column + 1
                    and offset in {cuft_offset, paperboard_offset, carton_price_offset, total_offset}
                    else None
                ),
                font_name="Times New Roman" if column in dimension_columns and offset in dimension_row_offsets else "宋体",
            )
    _apply_table_borders(
        sheet,
        packaging_start_row,
        packaging_start_row + len(packaging_rows) - 1,
        side_start_column,
        side_end_column,
    )

    function_start_row = packaging_start_row + len(packaging_rows) + 1
    function_end_row = function_start_row + 7
    _template_style_range(
        sheet,
        function_start_row,
        function_end_row,
        side_start_column,
        side_end_column,
    )
    sheet.merge_cells(
        start_row=function_start_row,
        start_column=side_start_column,
        end_row=function_end_row,
        end_column=side_end_column,
    )
    function_cell = sheet.cell(function_start_row, side_start_column)
    function_cell.value = f"功能介绍：{quote.remark or ''}"
    function_cell.font = Font(name="宋体", size=12, color=TEMPLATE_RED)
    function_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )
    _apply_table_borders(
        sheet,
        function_start_row,
        function_end_row,
        side_start_column,
        side_end_column,
    )

    # BuzzBee is the only customer whose internal workbook carries a separate
    # customer-facing color-box tier table. All other customers keep color-
    # box cost in the ordinary left-side "彩盒/内卡" detail row only.
    color_end_row = function_end_row
    if _is_buzzbee_customer(getattr(quote, "customer", "")):
        color_box = _dict_value(
            _dict_value(sales_payload.get("customer_quote_fields", {})).get("buzzbee", {})
        )
        color_tiers = _list_of_dicts(color_box.get("color_box_tiers", []))[:2]
        # Match the BuzzBee reference template: keep one completely blank row
        # between the function-introduction box and the quotation block.
        color_title_row = function_end_row + 2
        _template_style_range(
            sheet,
            color_title_row,
            color_title_row,
            side_start_column,
            side_start_column + 2,
        )
        sheet.merge_cells(
            start_row=color_title_row,
            start_column=side_start_column,
            end_row=color_title_row,
            end_column=side_start_column + 2,
        )
        _template_cell(
            sheet,
            color_title_row,
            side_start_column,
            "彩盒价格",
            wrap_text=False,
        )
        color_header_row = color_title_row + 1
        for column, label in enumerate(
            ("报客彩盒", "报客彩盒FSC", "MOQ数量"),
            start=side_start_column,
        ):
            _template_cell(sheet, color_header_row, column, label, wrap_text=False)
        for offset in range(2):
            row_index = color_header_row + 1 + offset
            tier = color_tiers[offset] if offset < len(color_tiers) else {}
            for column, value in enumerate(
                (
                    _number(tier.get("quote_price_hkd", "")),
                    _number(tier.get("fsc_price_hkd", "")),
                    _safe_text(tier.get("moq", "")),
                ),
                start=side_start_column,
            ):
                _template_cell(
                    sheet,
                    row_index,
                    column,
                    value,
                    color=TEMPLATE_RED,
                    wrap_text=False,
                    number_format=(
                        "0.00_);[Red]\\(0.00\\)"
                        if column in {side_start_column, side_start_column + 1}
                        else None
                    ),
                )
        color_end_row = color_header_row + 2
        _apply_table_borders(
            sheet,
            color_title_row,
            color_end_row,
            side_start_column,
            side_start_column + 2,
        )

    test_header_row = (
        color_end_row + 3
        if testing_fee_enabled and color_end_row > function_end_row
        else function_end_row + 2
        if testing_fee_enabled
        else None
    )
    testing_display_rows = (
        testing_fee_tiers or [
            {"moq": int(tier["moq"]), "unit_price_usd": 0.0}
            for tier in markup_tiers
        ]
        if testing_fee_enabled
        else []
    )
    if test_header_row is not None:
        _template_cell(
            sheet,
            test_header_row,
            side_start_column,
            "测试费用",
            wrap_text=False,
        )
        _template_cell(
            sheet,
            test_header_row,
            side_start_column + 1,
            testing_fee_total_usd,
            wrap_text=False,
            number_format='"US$"#,##0.00',
        )
        _template_cell(
            sheet,
            test_header_row,
            side_start_column + 2,
            "",
            wrap_text=False,
        )

    # Additional tax and Indonesia freight are direct costs.  They must remain
    # auditable as separate rows and must not be labelled as the percentage-
    # based miscellaneous charge calculated later from ``price * Q6``.  The
    # cost context has already applied the quote-region rule, so a mainland
    # quote reaches the exporter with an effective Indonesia freight of zero.
    direct_cost_rows = [
        ("附加税", _float_value(shipping.get("additional_tax_hkd"))),
        ("印尼运费", _float_value(cost_context.get("indonesia_freight_hkd"))),
    ]
    direct_cost_rows = [
        (label, amount)
        for label, amount in direct_cost_rows
        if amount > 0
    ]
    next_detail_row = detail_data_end_row + 1
    for label, amount in direct_cost_rows:
        for column, value in enumerate(("", label, label, amount), start=1):
            _template_cell(
                sheet,
                next_detail_row,
                column,
                value,
                horizontal="left" if column in {2, 3} else "center",
                wrap_text=False,
                number_format="0.000" if column == 4 else None,
                border=Border(left=TEMPLATE_MEDIUM) if column == 1 else None,
            )
        next_detail_row += 1

    detail_formula_end_row = next_detail_row - 1
    freight_row: int | None = None
    lifting_row: int | None = None
    if route_rows and (freight_output_enabled or lifting_output_enabled):
        route_header_row = next_detail_row
        detail_formula_end_row = route_header_row
        for column, value in enumerate(("", "运输方案", "运输方案", ""), start=1):
            _template_cell(
                sheet,
                route_header_row,
                column,
                value,
                horizontal="left" if column in {2, 3} else "center",
                wrap_text=False,
                border=Border(left=TEMPLATE_MEDIUM) if column == 1 else None,
            )
        for offset, route in enumerate(route_rows):
            _template_cell(
                sheet,
                route_header_row,
                route_start_column + offset,
                _safe_text(route.get("name") or route.get("item") or f"运输方案{offset + 1}"),
                wrap_text=False,
                border=None,
            )
        route_fee_rows: list[tuple[int, str, str]] = []
        next_route_fee_row = route_header_row + 1
        if freight_output_enabled:
            freight_row = next_route_fee_row
            route_fee_rows.append((freight_row, "运费", "freight_hkd"))
            next_route_fee_row += 1
        if lifting_output_enabled:
            lifting_row = next_route_fee_row
            route_fee_rows.append((lifting_row, "吊柜费", "lift_hkd"))
            next_route_fee_row += 1
        for row_index, label, value_key in route_fee_rows:
            _template_cell(
                sheet,
                row_index,
                2,
                label,
                horizontal="left",
                wrap_text=False,
                border=Border(left=TEMPLATE_MEDIUM),
            )
            _template_cell(
                sheet,
                row_index,
                3,
                label,
                horizontal="left",
                wrap_text=False,
                border=None,
            )
            _template_cell(sheet, row_index, 4, "", border=None)
            for offset, route in enumerate(route_rows):
                _template_cell(
                    sheet,
                    row_index,
                    route_start_column + offset,
                    _number(route.get(value_key, "")),
                    wrap_text=False,
                    number_format="0.00",
                    border=None,
                )

    rendered_route_fee_rows = [row for row in (freight_row, lifting_row) if row is not None]
    subtotal_row = (max(rendered_route_fee_rows) + 1) if rendered_route_fee_rows else next_detail_row
    price_columns = [4, *range(route_start_column, route_end_column + 1)]
    for column in price_columns:
        _template_cell(
            sheet,
            subtotal_row,
            column,
            "",
            bold=True,
            wrap_text=False,
            number_format="0.00",
            border=Border(bottom=TEMPLATE_THIN),
        )
    sheet.cell(subtotal_row, 4).value = f"=SUM(D{detail_start_row}:D{detail_formula_end_row})"
    if rendered_route_fee_rows:
        for column in range(route_start_column, route_end_column + 1):
            letter = get_column_letter(column)
            route_fee_terms = "+".join(f"{letter}{row}" for row in rendered_route_fee_rows)
            sheet.cell(subtotal_row, column).value = f"=$D${subtotal_row}+{route_fee_terms}"

    pricing_rows: list[dict[str, int | float]] = []
    detached_detail_ranges: list[tuple[int, int]] = []
    pricing_groups = _list_of_dicts(shipping.get("pricing_groups", []))
    pricing_entries = _list_of_dicts(shipping.get("pricing_entries", []))
    split_pricing = shipping.get("pricing_mode") == "component" or len(pricing_groups) > 1
    standard_detail_split = (
        shipping.get("pricing_mode") == "standard" and len(pricing_groups) > 1
    )
    # Ordinary quotes with detached detail multipliers keep the original
    # subtotal row as the main-multiplier cost row.  This is the established
    # management-sheet flow: main cost/×/÷/quote first, followed by each
    # detached detail block, and then the original summary tables.
    pricing_start_row = subtotal_row if standard_detail_split else subtotal_row + 1
    if split_pricing:
        def detached_group_entries(
            index: int,
            group: dict[str, Any],
        ) -> list[dict[str, Any]]:
            if not standard_detail_split or index == 0:
                return []
            group_name = _safe_text(group.get("name"))
            group_markup = _float_value(group.get("markup"))
            return [
                entry
                for entry in pricing_entries
                if _safe_text(entry.get("label")) == group_name
                and entry.get("markup_override") not in (None, "")
                and abs(
                    _float_value(entry.get("markup_override")) - group_markup
                )
                < 0.0000001
            ]

        group_layouts: list[dict[str, object]] = []
        next_group_row = pricing_start_row
        for index, group in enumerate(pricing_groups):
            entries = detached_group_entries(index, group)
            detail_row_start = next_group_row
            cost_row = detail_row_start + len(entries)
            markup_row = cost_row + 1
            settlement_row = cost_row + 2
            quote_row = cost_row + 3
            group_layouts.append(
                {
                    "index": index,
                    "group": group,
                    "entries": entries,
                    "detail_row_start": detail_row_start,
                    "cost_row": cost_row,
                    "markup_row": markup_row,
                    "settlement_row": settlement_row,
                    "quote_row": quote_row,
                }
            )
            next_group_row = quote_row + (
                2
                if standard_detail_split and index < len(pricing_groups) - 1
                else 1
            )

        group_quote_rows: list[int] = []
        first_markup_row = int(group_layouts[0]["markup_row"])
        for layout in group_layouts:
            index = int(layout["index"])
            group = layout["group"]
            assert isinstance(group, dict)
            entries = layout["entries"]
            assert isinstance(entries, list)
            detail_row_start = int(layout["detail_row_start"])
            cost_row = int(layout["cost_row"])
            markup_row = int(layout["markup_row"])
            settlement_row = int(layout["settlement_row"])
            quote_row = int(layout["quote_row"])
            group_quote_rows.append(quote_row)
            group_name = _safe_text(group.get("name")) or f"分项 {index + 1}"

            for entry_offset, entry in enumerate(entries):
                detail_row = detail_row_start + entry_offset
                category = _pricing_entry_summary_category(entry)
                description = _safe_text(entry.get("label")) or category
                source_row = next(
                    (
                        row
                        for row in range(detail_start_row, detail_formula_end_row + 1)
                        if _safe_text(sheet.cell(row, 2).value) == category
                        and _safe_text(sheet.cell(row, 3).value) == description
                    ),
                    None,
                )
                if source_row is None:
                    source_row = next(
                        (
                            row
                            for row in range(detail_start_row, detail_formula_end_row + 1)
                            if _safe_text(sheet.cell(row, 2).value) == category
                        ),
                        None,
                    )
                amount_numeric = _float_value(entry.get("amount_hkd"))
                if source_row is not None and amount_numeric > 0:
                    source_cell = sheet.cell(source_row, 4)
                    source_value = source_cell.value
                    amount_label = format(amount_numeric, ".10g")
                    if isinstance(source_value, str) and source_value.startswith("="):
                        source_cell.value = f"=({source_value[1:]})-{amount_label}"
                    else:
                        remaining = _float_value(source_value) - amount_numeric
                        source_cell.value = 0.0 if abs(remaining) < 0.0000001 else remaining
                amount_formula = _pricing_entry_amount(entry)
                if not (
                    isinstance(amount_formula, str) and amount_formula.startswith("=")
                ):
                    amount_formula = f"={format(_float_value(amount_formula), '.10g')}"
                for column, value in enumerate(
                    (
                        _pricing_entry_tax_tag(entry),
                        category,
                        description,
                        amount_formula,
                    ),
                    start=1,
                ):
                    _template_cell(
                        sheet,
                        detail_row,
                        column,
                        value,
                        horizontal="left" if column in {2, 3} else "center",
                        wrap_text=False,
                        number_format="0.000" if column == 4 else None,
                        border=(
                            Border(left=TEMPLATE_MEDIUM)
                            if column == 1 and not standard_detail_split
                            else None
                        ),
                    )
                for column in range(route_start_column, route_end_column + 1):
                    _template_cell(
                        sheet,
                        detail_row,
                        column,
                        f"=$D${detail_row}",
                        wrap_text=False,
                        number_format="0.000",
                        border=None,
                    )

            if entries:
                detached_detail_ranges.append(
                    (detail_row_start, detail_row_start + len(entries) - 1)
                )

            _template_cell(sheet, cost_row, 2, group_name, horizontal="left", wrap_text=False, border=None)
            _template_cell(sheet, cost_row, 3, "成本", horizontal="right", wrap_text=False, border=None)
            _template_cell(sheet, markup_row, 3, "×", horizontal="right", wrap_text=False, border=None)
            _template_cell(sheet, settlement_row, 3, "÷", horizontal="right", wrap_text=False, border=None)
            _template_cell(
                sheet,
                quote_row,
                2,
                f"{group_name}报价",
                horizontal="left",
                wrap_text=False,
                border=(
                    None
                    if standard_detail_split
                    else Border(bottom=TEMPLATE_THIN)
                ),
            )
            for column in price_columns:
                letter = get_column_letter(column)
                pricing_base = _float_value(
                    group.get("pricing_base_hkd", group.get("cost_hkd"))
                )
                if standard_detail_split and index == 0 and column == 4:
                    # The main multiplier reuses the original subtotal row.
                    # Preserve its SUM formula instead of replacing it with a
                    # circular self-reference such as ``=D38``.
                    cost_value: object = sheet.cell(cost_row, column).value
                elif standard_detail_split and index > 0 and entries and column == 4:
                    cost_value = (
                        f"=SUM(D{detail_row_start}:D{cost_row - 1})"
                    )
                elif standard_detail_split and column != 4:
                    cost_value = f"=$D${cost_row}"
                else:
                    cost_value = pricing_base
                if index == 0 and column != 4 and rendered_route_fee_rows:
                    route_fee_terms = "+".join(f"{letter}{row}" for row in rendered_route_fee_rows)
                    cost_value = f"=$D${cost_row}+{route_fee_terms}"
                _template_cell(sheet, cost_row, column, cost_value, wrap_text=False, number_format="0.000", border=None)
                _template_cell(
                    sheet,
                    markup_row,
                    column,
                    (
                        _float_value(group.get("markup"))
                        if column == 4
                        else f"=$D${markup_row}"
                    ),
                    wrap_text=False,
                    number_format="0.00",
                    border=None,
                )
                _template_cell(
                    sheet,
                    settlement_row,
                    column,
                    "=1-$Q$6" if column == 4 else f"=$D${settlement_row}",
                    wrap_text=False,
                    number_format="0.0000",
                    border=Border(bottom=TEMPLATE_THIN),
                )
                _template_cell(sheet, quote_row, column, f"={letter}{cost_row}*{letter}{markup_row}/{letter}{settlement_row}", bold=True, wrap_text=False, number_format="0.00", border=Border(bottom=TEMPLATE_THIN))
            if standard_detail_split:
                # Ordinary-customer multiplier segments use horizontal
                # separators only, matching the released management sheet.
                for column in price_columns:
                    _replace_border_sides(
                        sheet.cell(detail_row_start, column),
                        top=TEMPLATE_THIN,
                    )
            else:
                _apply_outline_border(
                    sheet,
                    detail_row_start,
                    quote_row,
                    2,
                    max(price_columns),
                    TEMPLATE_MEDIUM,
                )
        total_quote_row = next_group_row
        usd_row = total_quote_row + 1
        included_row = total_quote_row + 2 if testing_fee_enabled else usd_row
        _template_cell(
            sheet,
            total_quote_row,
            2,
            "报价合计",
            horizontal="left",
            wrap_text=False,
            bold=True,
            border=(
                None
                if standard_detail_split
                else Border(bottom=TEMPLATE_THIN)
            ),
        )
        if testing_fee_enabled:
            _template_cell(sheet, included_row, 2, "包含测试费用（US）：", horizontal="left", wrap_text=False, border=None)
        for column in price_columns:
            letter = get_column_letter(column)
            quote_terms = "+".join(f"{letter}{row}" for row in group_quote_rows) or "0"
            _template_cell(sheet, total_quote_row, column, f"={quote_terms}", bold=True, wrap_text=False, number_format="0.00", border=Border(bottom=TEMPLATE_THIN))
            _template_cell(sheet, usd_row, column, f"={letter}{total_quote_row}/$R$4", bold=True, wrap_text=False, number_format="0.00", border=None)
            if testing_fee_enabled:
                _template_cell(sheet, included_row, column, f"={letter}{usd_row}", bold=True, wrap_text=False, number_format="0.00", border=None)
        pricing_rows.append({
            "moq": int(active_tier["moq"]),
            "markup": float(active_tier["markup"]),
            "markup_row": first_markup_row,
            "settlement_row": first_markup_row + 1,
            "quote_row": total_quote_row,
            "usd_row": usd_row,
            "included_row": included_row,
        })
    else:
        pricing_row_stride = 6 if testing_fee_enabled else 5
        for index, tier in enumerate(markup_tiers):
            markup_row = pricing_start_row + index * pricing_row_stride
            settlement_row = markup_row + 1
            quote_row = markup_row + 2
            usd_row = markup_row + 3
            included_row = markup_row + 4 if testing_fee_enabled else usd_row
            pricing_rows.append({
                "moq": int(tier["moq"]),
                "markup": float(tier["markup"]),
                "markup_row": markup_row,
                "settlement_row": settlement_row,
                "quote_row": quote_row,
                "usd_row": usd_row,
                "included_row": included_row,
            })
            _template_cell(sheet, markup_row, 3, "×", border=None, horizontal="right")
            _template_cell(sheet, settlement_row, 3, "÷", border=None, horizontal="right")
            _template_cell(sheet, quote_row, 2, f"报价（{_moq_label(tier['moq'])}）", horizontal="left", wrap_text=False, border=Border(bottom=TEMPLATE_THIN))
            if testing_fee_enabled:
                _template_cell(sheet, included_row, 2, "包含测试费用（US）：", horizontal="left", wrap_text=False, border=None)
            for column in price_columns:
                letter = get_column_letter(column)
                _template_cell(sheet, markup_row, column, float(tier["markup"]), wrap_text=False, number_format="0.00", border=None)
                _template_cell(sheet, settlement_row, column, "=1-$Q$6", wrap_text=False, number_format="0.0000", border=Border(bottom=TEMPLATE_THIN))
                _template_cell(sheet, quote_row, column, f"={letter}{subtotal_row}*{letter}{markup_row}/{letter}{settlement_row}", bold=True, wrap_text=False, number_format="0.00", border=Border(bottom=TEMPLATE_THIN))
                _template_cell(sheet, usd_row, column, f"={letter}{quote_row}/$R$4", bold=True, wrap_text=False, number_format="0.00", border=None)
                if testing_fee_enabled:
                    _template_cell(sheet, included_row, column, f"={letter}{usd_row}", bold=True, wrap_text=False, number_format="0.00", border=None)
    pricing_end_row = pricing_rows[-1]["included_row"]
    pricing_by_moq = {int(row["moq"]): row for row in pricing_rows}
    active_pricing = pricing_by_moq.get(int(active_tier["moq"]), pricing_rows[-1])

    test_adjusted_cells: dict[int, str] = {}
    for offset, test_tier in enumerate(testing_display_rows, start=1):
        assert test_header_row is not None
        row_index = test_header_row + offset
        moq = int(test_tier["moq"])
        _template_cell(
            sheet,
            row_index,
            side_start_column,
            moq,
            wrap_text=False,
            number_format="#,##0",
        )
        _template_cell(
            sheet,
            row_index,
            side_start_column + 1,
            _float_value(test_tier.get("unit_price_usd")),
            wrap_text=False,
            number_format='"US$"0.00',
        )
        pricing = pricing_by_moq.get(moq)
        if pricing is not None:
            adjusted_formula = (
                f"={get_column_letter(side_start_column + 1)}{row_index}"
                f"*D{pricing['markup_row']}"
            )
        else:
            adjusted_formula = (
                f"={get_column_letter(side_start_column + 1)}{row_index}"
                f"*D{active_pricing['markup_row']}"
            )
        _template_cell(
            sheet,
            row_index,
            side_start_column + 2,
            adjusted_formula,
            wrap_text=False,
            number_format='"US$"0.00',
        )
        test_adjusted_cells[moq] = (
            f"${get_column_letter(side_start_column + 2)}${row_index}"
        )
    test_end_row = function_end_row
    if test_header_row is not None:
        test_end_row = test_header_row + len(testing_display_rows)
        _apply_table_borders(
            sheet,
            test_header_row,
            test_end_row,
            side_start_column,
            side_start_column + 2,
        )
    for pricing in pricing_rows:
        adjusted_reference = test_adjusted_cells.get(int(pricing["moq"]))
        if adjusted_reference is None:
            continue
        for column in price_columns:
            letter = get_column_letter(column)
            sheet.cell(int(pricing["included_row"]), column).value = (
                f"={letter}{pricing['usd_row']}+{adjusted_reference}"
            )

    adaptive_content_end_row = max(
        int(pricing_end_row),
        test_end_row,
        color_end_row,
        function_end_row,
    )

    # C:P — formula-driven summary, anchored below whichever of the adaptive
    # left/right regions is longer. Detached multiplier details remain in the
    # original row flow and are aggregated directly by the summary formulas.
    summary_start_row = adaptive_content_end_row + 2
    amount_format = "0.00_);[Red]\\(0.00\\)"
    summary_detail_ranges = [
        (detail_start_row, detail_formula_end_row),
        *detached_detail_ranges,
    ]

    def sum_category(label_cell: str) -> str:
        return "=" + "+".join(
            f"SUMIF($B${start}:$B${end},{label_cell},$D${start}:$D${end})"
            for start, end in summary_detail_ranges
        )

    def sum_description(description: str) -> str:
        return "=" + "+".join(
            f'SUMIF($C${start}:$C${end},"{description}",$D${start}:$D${end})'
            for start, end in summary_detail_ranges
        )

    def sum_tax_tag(tax_tag: str) -> str:
        return "=" + "+".join(
            f'SUMIF($A${start}:$A${end},"{tax_tag}",$D${start}:$D${end})'
            for start, end in summary_detail_ranges
        )

    def style_summary_pair(
        header_row: int,
        headers: list[object],
        values: list[object],
        *,
        first_header_fill: str | None = None,
    ) -> None:
        for column in range(3, 17):
            index = column - 3
            header = headers[index] if index < len(headers) else ""
            value = values[index] if index < len(values) else ""
            fill = None
            if column == 3:
                fill = first_header_fill
            elif column <= 14:
                fill = "FFFF00"
            _template_cell(
                sheet,
                header_row,
                column,
                header,
                bold=column == 4,
                color=TEMPLATE_BLACK,
                fill=fill,
                size=10,
                wrap_text=False,
                number_format=amount_format,
            )
            _template_cell(
                sheet,
                header_row + 1,
                column,
                value,
                bold=column == 4,
                color=TEMPLATE_BLACK,
                size=10,
                wrap_text=False,
                number_format=amount_format,
            )

    first_header_row = summary_start_row
    first_value_row = first_header_row + 1
    style_summary_pair(
        first_header_row,
        ["旺季价", "货价", "料价", "吹气", "搪胶", "车发", "车衣", "五金", "电子", "马达", "吸塑", "彩盒/内卡"],
        [
            "按出厂货价核",
            f"=D{active_pricing['quote_row']}",
            f"=K{mold_total_row}" if mold_lines else sum_category(f"E{first_header_row}"),
            sum_category(f"F{first_header_row}"),
            sum_category(f"G{first_header_row}"),
            sum_category(f"H{first_header_row}"),
            sum_category(f"I{first_header_row}"),
            sum_category(f"J{first_header_row}"),
            sum_category(f"K{first_header_row}"),
            sum_category(f"L{first_header_row}"),
            sum_category(f"M{first_header_row}"),
            sum_category(f"N{first_header_row}"),
        ],
        first_header_fill="FFC000",
    )

    second_header_row = summary_start_row + 3
    second_value_row = second_header_row + 1
    third_header_row = summary_start_row + 6
    third_value_row = third_header_row + 1
    tax_header_row = summary_start_row + 9
    deduction_row = tax_header_row + 3
    style_summary_pair(
        second_header_row,
        ["ABS料价成本", "未减税前码数", "减税后码数", "电池", "利宝/说明书", "电镀", "其他外购", "纸箱", "运费", "吊柜费", "杂项", "不含人工成本"],
        [
            f'=SUMIF($D${mold_start_row}:$D${mold_end_row},"*ABS*",$K${mold_start_row}:$K${mold_end_row})',
            f"=IFERROR(D{first_value_row}/N{third_value_row},0)",
            f"=IFERROR(D{first_value_row}/P{deduction_row},0)",
            sum_category(f"F{second_header_row}"),
            sum_category(f"G{second_header_row}"),
            sum_category(f"H{second_header_row}"),
            sum_category(f"I{second_header_row}"),
            sum_category(f"J{second_header_row}"),
            0.0,
            0.0,
            f"=D{first_value_row}*$Q$6",
            (
                f"=SUM(E{first_value_row}:N{first_value_row},"
                f"F{second_value_row}:M{second_value_row},G{third_value_row})"
                f"+({sum_description('附加税')[1:]})"
                f"+({sum_description('印尼运费')[1:]})"
            ),
        ],
    )
    sheet.cell(second_header_row, 3).fill = PatternFill(fill_type=None)
    sheet.cell(second_header_row, 4).font = Font(name="宋体", size=10, bold=True, color=TEMPLATE_BLUE)
    sheet.cell(second_header_row, 5).font = Font(name="宋体", size=10, bold=True, color=TEMPLATE_RED)
    sheet.cell(second_value_row, 4).font = Font(name="宋体", size=10, bold=True, color=TEMPLATE_BLUE)
    sheet.cell(second_value_row, 5).font = Font(name="宋体", size=10, bold=True, color=TEMPLATE_RED)

    style_summary_pair(
        third_header_row,
        ["ABS占货价%", "", "啤工", "喷油工", "油漆", "装配工", "人工比例", "毛利", "毛利率", "利润", "利润率", "总成本"],
        [
            f"=IFERROR(C{second_value_row}/D{first_value_row},0)",
            "",
            f"=J{mold_total_row}" if mold_lines else sum_category(f"E{third_header_row}"),
            sum_category(f"F{third_header_row}"),
            sum_category(f"G{third_header_row}"),
            sum_category(f"H{third_header_row}"),
            f"=IFERROR((E{third_value_row}+F{third_value_row}+H{third_value_row})/D{first_value_row},0)",
            f"=D{first_value_row}-N{second_value_row}",
            f"=IFERROR(J{third_value_row}/D{first_value_row},0)",
            f"=D{first_value_row}-N{third_value_row}",
            f"=IFERROR(L{third_value_row}/D{first_value_row},0)",
            f"=N{second_value_row}+E{third_value_row}+F{third_value_row}+H{third_value_row}",
        ],
    )
    sheet.cell(third_header_row, 3).fill = PatternFill(fill_type=None)
    sheet.cell(third_value_row, 3).number_format = "0.0%"
    for column in (9, 11, 13):
        sheet.cell(third_value_row, column).number_format = "0.0%"
    sheet.cell(third_header_row, 14).font = Font(name="宋体", size=10, bold=True, color=TEMPLATE_BLUE)
    sheet.cell(third_value_row, 14).font = Font(name="宋体", size=10, bold=True, color=TEMPLATE_BLUE)

    tax_keys = (
        "tax1",
        "slush3",
        "sewhair13",
        "sewcloth13",
        "suction6",
        "freight9",
        "tax13b",
        "carton",
        "labor13",
    )
    tax_headers = ["人民币外购件成本", "含税13%类成本", ""] + [
        _safe_text(t4.get(key, {}).get("label", "")) for key in tax_keys
    ]
    for column in range(3, 17):
        value = tax_headers[column - 3] if column - 3 < len(tax_headers) else ""
        _template_cell(
            sheet,
            tax_header_row,
            column,
            value,
            color=TEMPLATE_BLACK,
            fill=TEMPLATE_SKY if 6 <= column <= 14 else None,
            size=10,
            wrap_text=False,
        )

    tax_amount_row = tax_header_row + 1
    tax_13_formula = sum_tax_tag("¥13%")
    rmb_purchase_formula = f"=D{tax_amount_row}+({sum_tax_tag('¥1%')[1:]})"
    for column in range(3, 17):
        if column == 3:
            value: object = rmb_purchase_formula
            number_format = amount_format
        elif column == 4:
            value = tax_13_formula
            number_format = amount_format
        elif column == 5:
            value = ""
            number_format = amount_format
        elif 6 <= column <= 14:
            key = tax_keys[column - 6]
            rate_percent = t4.get(key, {}).get("rate_percent")
            value = (
                ""
                if key == "carton" or (key == "sewcloth13" and rate_percent in (None, ""))
                else _percent_fraction(rate_percent)
            )
            number_format = "0.00%"
        elif column == 15:
            value = "合计减税"
            number_format = amount_format
        else:
            value = "减税后成本"
            number_format = amount_format
        _template_cell(
            sheet,
            tax_amount_row,
            column,
            value,
            bold=column in {15, 16},
            color=TEMPLATE_RED if column in {15, 16} else TEMPLATE_BLACK,
            fill=TEMPLATE_SKY if 6 <= column <= 14 else None,
            size=9 if column in {15, 16} else 10,
            wrap_text=False,
            number_format=number_format,
        )

    for column in range(3, 17):
        _template_cell(
            sheet,
            deduction_row,
            column,
            "",
            bold=column in {15, 16},
            color=TEMPLATE_RED if column in {15, 16} else TEMPLATE_BLACK,
            fill=TEMPLATE_SKY if 6 <= column <= 14 else None,
            size=9 if column in {15, 16} else 10,
            wrap_text=False,
            number_format=amount_format,
        )
    tax_amount_references = {
        "tax1": f"H{second_value_row}",
        "slush3": f"G{first_value_row}",
        "sewhair13": f"H{first_value_row}",
        "suction6": f"M{first_value_row}",
        "freight9": f"K{second_value_row}",
        "tax13b": f"D{tax_amount_row}",
        "carton": f"J{second_value_row}",
        "labor13": f"SUM(E{third_value_row}:F{third_value_row},H{third_value_row})",
    }
    for column, key in enumerate(tax_keys, start=6):
        rate_percent = t4.get(key, {}).get("rate_percent")
        if key == "carton" or (key == "sewcloth13" and rate_percent in (None, "")):
            sheet.cell(deduction_row, column).value = ""
        else:
            amount_reference = (
                format(_float_value(t4.get(key, {}).get("amount_hkd")), ".10g")
                if key == "sewcloth13"
                else tax_amount_references[key]
            )
            sheet.cell(deduction_row, column).value = (
                f"={amount_reference}*{get_column_letter(column)}{tax_amount_row}"
            )
    sheet.cell(deduction_row, 15).value = f"=SUM(F{deduction_row}:N{deduction_row})"
    sheet.cell(deduction_row, 16).value = f"=N{third_value_row}-O{deduction_row}"

    # Workbook presentation/printing settings.
    widths = {
        "A": 3.5833333333,
        "B": 7.5,
        "C": 26.5,
        "D": 10.3333333333,
        "E": 12.625,
        "F": 16.75,
        "G": 13,
        "H": 13,
        "I": 13,
        "J": 13,
        "K": 11.0833333333,
        "L": 13,
        "M": 13,
        "N": 10.3333333333,
        "O": 8.25,
        "P": 13,
        "Q": 13,
        "R": 13,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for column in range(19, outer_right_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 13
    _add_product_image_to_region(
        sheet,
        attachments,
        min_row=8,
        max_row=packaging_start_row - 1,
        min_column=side_start_column,
        max_column=side_end_column,
    )
    for row in range(summary_start_row, deduction_row + 1):
        sheet.row_dimensions[row].height = 15
    for row in (summary_start_row + 2, summary_start_row + 5, summary_start_row + 8):
        sheet.row_dimensions[row].height = 7
    sheet.row_dimensions[tax_header_row + 2].height = 6
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    sheet.freeze_panes = "A8"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.35
    sheet.page_margins.bottom = 0.35
    _apply_outline_border(
        sheet,
        1,
        deduction_row,
        1,
        outer_right_column,
        TEMPLATE_MEDIUM,
    )
    _apply_outline_border(
        sheet,
        7,
        7,
        1,
        outer_right_column,
        TEMPLATE_MEDIUM,
    )
    sheet.print_area = (
        f"A1:{get_column_letter(outer_right_column)}{deduction_row}"
    )
    sheet.sheet_properties.outlinePr.summaryBelow = True
    if shipping.get("pricing_mode") == "component" and pricing_groups:
        _replace_with_component_summary_sheet(
            workbook,
            quote,
            sections,
            rr2_cost_summary,
            attachments,
            side_start_column=side_start_column,
            side_end_column=side_end_column,
            outer_right_column=outer_right_column,
        )


def _replace_with_component_summary_sheet(
    workbook: Workbook,
    quote: InternalQuote,
    sections: list[InternalQuoteSection],
    rr2_cost_summary: dict[str, Any],
    attachments: list[InternalQuoteAttachment] | None,
    *,
    side_start_column: int,
    side_end_column: int,
    outer_right_column: int,
) -> None:
    """Replace the compact JustPlay view with the released component layout.

    The original formula sheet is retained as a hidden calculation source for
    the unchanged carton, testing-fee and final cost-summary blocks.  The first
    visible sheet then mirrors the business workbook: each component owns one
    complete molding/cost/pricing segment, followed by one packaging and
    transport segment and one final summary.
    """

    source = workbook["报价明细"]
    source.title = "_报价明细计算"
    target = workbook.create_sheet("报价明细", 0)
    shipping = _dict_value(rr2_cost_summary.get("shipping_pricing", {}))
    pricing_groups = _list_of_dicts(shipping.get("pricing_groups", []))
    pricing_entries = _list_of_dicts(shipping.get("pricing_entries", []))
    global_pricing = _dict_value(shipping.get("global_pricing", {}))
    global_entries = _list_of_dicts(global_pricing.get("entries", []))
    by_code = {section.department: section for section in sections}
    source_summary_start = next(
        (
            row
            for row in range(1, source.max_row + 1)
            if str(source.cell(row, 3).value or "") == "旺季价"
        ),
        source.max_row + 1,
    )
    source_packaging_start = next(
        (
            row
            for row in range(8, source_summary_start)
            if str(source.cell(row, side_start_column).value or "").startswith(
                ("外箱", "内箱", "彩盒尺寸", "产品尺寸")
            )
        ),
        source_summary_start,
    )
    source_packaging_total_row = next(
        (
            row
            for row in range(source_packaging_start, source_summary_start)
            if str(source.cell(row, side_start_column).value or "") == "合计"
        ),
        None,
    )

    def copy_cell(source_cell, target_cell, *, link_computed: bool = False) -> None:
        value = source_cell.value
        if link_computed and value not in (None, ""):
            if isinstance(value, str) and not value.startswith("="):
                target_cell.value = value
            else:
                target_cell.value = (
                    f"={quote_sheetname(source.title)}!{source_cell.coordinate}"
                )
        else:
            target_cell.value = value
        if source_cell.has_style:
            target_cell.font = copy(source_cell.font)
            target_cell.fill = copy(source_cell.fill)
            target_cell.border = copy(source_cell.border)
            target_cell.alignment = copy(source_cell.alignment)
            target_cell.protection = copy(source_cell.protection)
            target_cell.number_format = source_cell.number_format
        if source_cell.hyperlink is not None:
            target_cell._hyperlink = copy(source_cell.hyperlink)
        if source_cell.comment is not None:
            target_cell.comment = copy(source_cell.comment)

    def copy_block(
        min_row: int,
        max_row: int,
        min_column: int,
        max_column: int,
        destination_row: int,
        *,
        link_computed: bool = False,
    ) -> int:
        row_offset = destination_row - min_row
        for source_row in range(min_row, max_row + 1):
            target_row = source_row + row_offset
            target.row_dimensions[target_row].height = source.row_dimensions[source_row].height
            for column in range(min_column, max_column + 1):
                copy_cell(
                    source.cell(source_row, column),
                    target.cell(target_row, column),
                    link_computed=link_computed,
                )
        for merged_range in list(source.merged_cells.ranges):
            if (
                merged_range.min_row >= min_row
                and merged_range.max_row <= max_row
                and merged_range.min_col >= min_column
                and merged_range.max_col <= max_column
            ):
                target.merge_cells(
                    start_row=merged_range.min_row + row_offset,
                    start_column=merged_range.min_col,
                    end_row=merged_range.max_row + row_offset,
                    end_column=merged_range.max_col,
                )
        return max_row + row_offset

    for column in range(1, outer_right_column + 1):
        letter = get_column_letter(column)
        target.column_dimensions[letter].width = source.column_dimensions[letter].width
        target.column_dimensions[letter].hidden = source.column_dimensions[letter].hidden
    copy_block(1, 7, 1, outer_right_column, 1)
    material_reference_cells: list[tuple[str, float, str]] = []
    machine_reference_cells: list[tuple[str, float, str]] = []
    for header_row, start_column, end_column in (
        (1, 4, 13),
        (3, 4, 11),
        (5, 4, 13),
    ):
        candidates = (
            machine_reference_cells if header_row == 5 else material_reference_cells
        )
        for column in range(start_column, end_column + 1):
            label = _safe_text(source.cell(header_row, column).value)
            price = _float_value(source.cell(header_row + 1, column).value)
            if label and price > 0:
                candidates.append(
                    (
                        label,
                        price,
                        f"{get_column_letter(column)}${header_row + 1}",
                    )
                )

    mold_headers = (
        "",
        "",
        "名称",
        "料型",
        "料重(G)",
        "料价(G)",
        "机型",
        "1出几套",
        "目标数",
        "啤工",
        "料金额",
        "报价啤工",
    )
    mold_lines = _calculation_lines(by_code.get("molding"), "injection")
    if not mold_lines:
        mold_lines = _list_of_dicts(
            _section_payload(by_code.get("engineering")).get("molds", [])
        )
    component_ids = {str(group.get("id") or "") for group in pricing_groups}
    default_component_id = str(pricing_groups[0].get("id") or "")

    def assigned_component_id(row: dict[str, Any]) -> str:
        component_id = str(row.get("pricing_component_id") or "")
        return component_id if component_id in component_ids else default_component_id

    component_usd_rows: list[int] = []
    component_mold_ranges: list[tuple[int, int]] = []
    visible_detail_ranges: list[tuple[int, int]] = []
    current_row = 8
    for index, group in enumerate(pricing_groups):
        component_id = str(group.get("id") or "")
        component_name = _safe_text(group.get("name")) or f"配件 {index + 1}"
        component_molds = [
            line for line in mold_lines if assigned_component_id(line) == component_id
        ]
        component_entries = [
            entry
            for entry in pricing_entries
            if not bool(entry.get("is_global"))
            and assigned_component_id(entry) == component_id
            and str(entry.get("kind") or "") != "injection"
        ]

        title_row = current_row
        target.merge_cells(
            start_row=title_row,
            start_column=1,
            end_row=title_row,
            end_column=12,
        )
        _template_cell(
            target,
            title_row,
            1,
            f"{component_name}-明细",
            color=TEMPLATE_BLUE,
            fill="FFFF00",
            bold=True,
            wrap_text=False,
        )
        target.cell(title_row, 1).alignment = Alignment(horizontal="center", vertical="center")
        target.row_dimensions[title_row].height = 18

        header_row = title_row + 1
        for column, header in enumerate(mold_headers, start=1):
            _template_cell(
                target,
                header_row,
                column,
                header,
                color=TEMPLATE_RED if column not in {1, 2} else TEMPLATE_BLACK,
                wrap_text=False,
            )
        mold_start_row = header_row + 1
        mold_slots = max(6, len(component_molds))
        mold_end_row = mold_start_row + mold_slots - 1
        for mold_index in range(mold_slots):
            row_index = mold_start_row + mold_index
            line = component_molds[mold_index] if mold_index < len(component_molds) else {}
            material_price_g = _float_value(line.get("material_price_hkd_g"))
            if material_price_g <= 0:
                material_price_g = _float_value(line.get("material_price_hkd_lb")) / 454
            values = (
                mold_index + 1 if line else "",
                "",
                _safe_text(line.get("item") or line.get("chinese_name") or line.get("mold_no") or ""),
                _safe_text(line.get("material") or line.get("material_type") or ""),
                _number(line.get("loss_weight_g", line.get("net_weight_g", ""))),
                material_price_g if line and material_price_g > 0 else "",
                _safe_text(line.get("machine_code") or line.get("machine_name") or ""),
                _number(line.get("sets", line.get("cavity", ""))),
                _number(line.get("target_output", "")),
                _number(line.get("molding_cost_hkd", "")),
                _number(line.get("material_cost_hkd", "")),
                _number(line.get("customer_price_hkd", "")),
            )
            for column, value in enumerate(values, start=1):
                _template_cell(
                    target,
                    row_index,
                    column,
                    value,
                    horizontal="left" if column in {2, 3, 4} else "center",
                    wrap_text=False,
                    number_format=(
                        "0.0000_ " if column in {6, 10, 12} else "0.000_ " if column == 11 else None
                    ),
                    font_name="Times New Roman" if column in {1, 5, 6, 8, 9, 10, 11, 12} else "宋体",
                )
            if line:
                _apply_molding_row_formulas(
                    target,
                    row_index,
                    line,
                    material_reference_cells,
                    machine_reference_cells,
                )
            if line and values[11] == "":
                target.cell(row_index, 12).value = f"=J{row_index}*1.15"
                target.cell(row_index, 12).number_format = "0.0000_ "
            target.row_dimensions[row_index].height = 15

        mold_total_row = mold_end_row + 1
        for column in range(1, 13):
            _template_cell(target, mold_total_row, column, "", color=TEMPLATE_RED, wrap_text=False)
        target.cell(mold_total_row, 3).value = "模具合计"
        target.cell(mold_total_row, 3).font = Font(
            name="宋体", size=10, bold=True, color=TEMPLATE_BLUE
        )
        for column in (5, 10, 11, 12):
            letter = get_column_letter(column)
            target.cell(mold_total_row, column).value = (
                f"=SUM({letter}{mold_start_row}:{letter}{mold_end_row})"
            )
            target.cell(mold_total_row, column).number_format = (
                "0.0000_ " if column in {10, 12} else "0.000_ "
            )
        _apply_table_borders(target, header_row, mold_total_row, 1, 12)
        component_mold_ranges.append((mold_start_row, mold_end_row))

        detail_title_row = mold_total_row + 3
        _template_cell(
            target,
            detail_title_row,
            3,
            f"{component_name}明细",
            color=TEMPLATE_BLUE,
            fill="FFFF00",
            bold=True,
            horizontal="left",
            wrap_text=False,
        )
        _template_cell(target, detail_title_row, 4, "出厂价", wrap_text=False)
        detail_row = detail_title_row + 1
        if component_molds:
            detail_values: list[tuple[str, str, str, object]] = [
                ("¥13%", "料价", "料价", f"=K{mold_total_row}"),
                ("", "啤工", "啤工", f"=J{mold_total_row}"),
            ]
        else:
            detail_values = []
        for entry in component_entries:
            category = _pricing_entry_summary_category(entry)
            detail_values.append(
                (
                    _pricing_entry_tax_tag(entry),
                    category,
                    _safe_text(entry.get("label")) or category,
                    _pricing_entry_amount(entry),
                )
            )
        expected_cost = _float_value(group.get("cost_hkd"))
        mold_cost = sum(
            _float_value(line.get("material_cost_hkd"))
            + _float_value(line.get("molding_cost_hkd"))
            for line in component_molds
        )
        listed_cost = mold_cost + sum(
            _float_value(entry.get("amount_hkd")) for entry in component_entries
        )
        residual = expected_cost - listed_cost
        if abs(residual) >= 0.0005:
            detail_values.append(("", "其他成本", "分配调整", residual))
        if not detail_values:
            detail_values.append(("", "其他成本", "配件成本", expected_cost))

        for tax_tag, category, description, amount in detail_values:
            for column, value in enumerate((tax_tag, category, description, amount), start=1):
                _template_cell(
                    target,
                    detail_row,
                    column,
                    value,
                    horizontal="left" if column in {2, 3} else "center",
                    wrap_text=False,
                    number_format="0.000" if column == 4 else None,
                    border=Border(left=TEMPLATE_MEDIUM) if column == 1 else None,
                )
            detail_row += 1

        subtotal_row = detail_row
        visible_detail_ranges.append((detail_title_row + 1, subtotal_row - 1))
        _template_cell(target, subtotal_row, 3, "成本金额：", bold=True, horizontal="right", border=None)
        _template_cell(
            target,
            subtotal_row,
            4,
            f"=SUM(D{detail_title_row + 1}:D{subtotal_row - 1})",
            bold=True,
            number_format="0.00",
            border=None,
        )
        markup_row = subtotal_row + 1
        settlement_row = subtotal_row + 2
        hkd_row = subtotal_row + 3
        _template_cell(target, markup_row, 3, "×", horizontal="right", border=None)
        _template_cell(target, markup_row, 4, _float_value(group.get("markup")), number_format="0.000", border=None)
        _template_cell(target, settlement_row, 3, "÷", horizontal="right", border=Border(bottom=TEMPLATE_MEDIUM))
        _template_cell(target, settlement_row, 4, "=1-$Q$6", number_format="0.0000", border=Border(bottom=TEMPLATE_MEDIUM))
        _template_cell(target, hkd_row, 3, "配件报价（HKD）：", bold=True, color=TEMPLATE_BLUE, horizontal="right", border=None)
        _template_cell(target, hkd_row, 4, f"=D{subtotal_row}*D{markup_row}/D{settlement_row}", bold=True, color=TEMPLATE_BLUE, number_format="0.00", border=None)
        fx_row = hkd_row + 1
        usd_row = hkd_row + 2
        _template_cell(
            target,
            fx_row,
            3,
            "美金兑港币汇率：",
            color=TEMPLATE_BLUE,
            horizontal="right",
            border=None,
        )
        _template_cell(
            target,
            fx_row,
            4,
            "=$R$4",
            color=TEMPLATE_BLUE,
            number_format="0.00",
            border=None,
        )
        _template_cell(
            target,
            usd_row,
            3,
            "配件单价（USD）：",
            bold=True,
            color=TEMPLATE_BLUE,
            horizontal="right",
            border=None,
        )
        _template_cell(
            target,
            usd_row,
            4,
            f"=D{hkd_row}/D{fx_row}",
            bold=True,
            color=TEMPLATE_BLUE,
            number_format="0.000",
            border=None,
        )
        _apply_outline_border(target, detail_title_row, usd_row, 1, 8, TEMPLATE_MEDIUM)
        component_usd_rows.append(usd_row)
        current_row = usd_row + 3

    all_route_rows = _list_of_dicts(shipping.get("rows", []))
    if all_route_rows and _safe_text(all_route_rows[0].get("name")) == "出厂价":
        route_rows = all_route_rows[1:]
    else:
        route_rows = all_route_rows
    route_rows = route_rows[: max(0, outer_right_column - 4)]
    quote_price_columns = [4, *range(5, 5 + len(route_rows))]
    quote_price_end_column = max(quote_price_columns)

    packaging_title_row = current_row
    _template_cell(target, packaging_title_row, 3, "包装明细", fill="FFFF00", bold=True, color=TEMPLATE_BLUE, wrap_text=False)
    _template_cell(target, packaging_title_row, 4, "出厂价", fill="FFFF00", bold=True, color=TEMPLATE_BLUE, wrap_text=False)
    for column, route in enumerate(route_rows, start=5):
        _template_cell(
            target,
            packaging_title_row,
            column,
            _safe_text(route.get("name") or route.get("item"))
            or f"运输方案{column - 4}",
            fill="FFFF00",
            bold=True,
            color=TEMPLATE_BLUE,
            wrap_text=False,
        )
    packaging_detail_row = packaging_title_row + 1
    target_packaging_total_row = (
        packaging_title_row + source_packaging_total_row - source_packaging_start
        if source_packaging_total_row is not None
        and source_packaging_start < source_summary_start
        else None
    )
    carton_rendered = False
    for entry in global_entries:
        category = _pricing_entry_summary_category(entry)
        is_carton = category == "纸箱" or str(entry.get("kind") or "") == "carton"
        if is_carton and carton_rendered:
            continue
        description = _safe_text(entry.get("label")) or category
        if is_carton:
            carton_rendered = True
            description = "纸箱"
            amount: object = (
                f"={get_column_letter(side_start_column + 1)}{target_packaging_total_row}"
                if target_packaging_total_row is not None
                else _pricing_entry_amount(entry)
            )
        else:
            source_detail_row = next(
                (
                    row
                    for row in range(8, source_summary_start)
                    if _safe_text(source.cell(row, 2).value) == category
                    and _safe_text(source.cell(row, 3).value) == description
                ),
                None,
            )
            amount = (
                f"={quote_sheetname(source.title)}!D{source_detail_row}"
                if source_detail_row is not None
                else _pricing_entry_amount(entry)
            )
        for column, value in enumerate(
            (
                _pricing_entry_tax_tag(entry),
                category,
                description,
                amount,
            ),
            start=1,
        ):
            _template_cell(
                target,
                packaging_detail_row,
                column,
                value,
                horizontal="left" if column in {2, 3} else "center",
                wrap_text=False,
                number_format="0.000" if column == 4 else None,
                border=Border(left=TEMPLATE_MEDIUM) if column == 1 else None,
            )
        packaging_detail_row += 1
    if not global_entries:
        for column, value in enumerate(("", "包装", "包装及纸箱", _float_value(global_pricing.get("cost_hkd"))), start=1):
            _template_cell(target, packaging_detail_row, column, value, number_format="0.000" if column == 4 else None, border=None)
        packaging_detail_row += 1

    packaging_cost_end_row = packaging_detail_row - 1
    visible_detail_ranges.append((packaging_title_row + 1, packaging_cost_end_row))
    packaging_freight_row = packaging_detail_row
    _template_cell(
        target,
        packaging_freight_row,
        3,
        "运费",
        horizontal="left",
        wrap_text=False,
        border=None,
    )
    for column, route in enumerate(route_rows, start=5):
        _template_cell(
            target,
            packaging_freight_row,
            column,
            _float_value(route.get("freight_hkd")),
            number_format="0.000",
            border=None,
        )

    packaging_subtotal_row = packaging_freight_row + 1
    _template_cell(target, packaging_subtotal_row, 3, "成本金额：", bold=True, horizontal="right", border=None)
    for column in quote_price_columns:
        letter = get_column_letter(column)
        formula = f"=SUM(D{packaging_title_row + 1}:D{packaging_cost_end_row})"
        if column != 4:
            formula = f"=$D${packaging_subtotal_row}+{letter}{packaging_freight_row}"
        _template_cell(
            target,
            packaging_subtotal_row,
            column,
            formula,
            bold=True,
            number_format="0.00",
            border=None,
        )
    packaging_markup_row = packaging_subtotal_row + 1
    packaging_settlement_row = packaging_subtotal_row + 2
    packaging_hkd_row = packaging_subtotal_row + 3
    packaging_fx_row = packaging_subtotal_row + 4
    packaging_usd_row = packaging_subtotal_row + 5
    _template_cell(target, packaging_markup_row, 3, "×", horizontal="right", border=None)
    _template_cell(target, packaging_settlement_row, 3, "÷", horizontal="right", border=Border(bottom=TEMPLATE_MEDIUM))
    _template_cell(target, packaging_hkd_row, 3, "单价（HK$）：", bold=True, color=TEMPLATE_BLUE, horizontal="right", border=None)
    _template_cell(target, packaging_fx_row, 3, "美金兑港币汇率：", color=TEMPLATE_BLUE, horizontal="right", border=None)
    _template_cell(target, packaging_usd_row, 3, "包装成本（USD）：", bold=True, color=TEMPLATE_BLUE, horizontal="right", border=None)
    packaging_markup = _float_value(
        global_pricing.get("markup"),
        _float_value(pricing_groups[0].get("markup")),
    )
    for column in quote_price_columns:
        letter = get_column_letter(column)
        _template_cell(
            target,
            packaging_markup_row,
            column,
            packaging_markup if column == 4 else f"=$D${packaging_markup_row}",
            number_format="0.000",
            border=None,
        )
        _template_cell(
            target,
            packaging_settlement_row,
            column,
            "=1-$Q$6" if column == 4 else f"=$D${packaging_settlement_row}",
            number_format="0.0000",
            border=Border(bottom=TEMPLATE_MEDIUM),
        )
        _template_cell(
            target,
            packaging_hkd_row,
            column,
            f"={letter}{packaging_subtotal_row}*{letter}{packaging_markup_row}/{letter}{packaging_settlement_row}",
            bold=True,
            color=TEMPLATE_BLUE,
            number_format="0.00",
            border=None,
        )
        _template_cell(
            target,
            packaging_fx_row,
            column,
            "=$R$4",
            color=TEMPLATE_BLUE,
            number_format="0.00",
            border=None,
        )
        _template_cell(
            target,
            packaging_usd_row,
            column,
            f"={letter}{packaging_hkd_row}/{letter}{packaging_fx_row}",
            bold=True,
            color=TEMPLATE_BLUE,
            number_format="0.000",
            border=None,
        )
    _apply_outline_border(
        target,
        packaging_title_row,
        packaging_usd_row,
        1,
        max(8, quote_price_end_column),
        TEMPLATE_MEDIUM,
    )

    side_end_row = packaging_title_row
    if source_packaging_start < source_summary_start:
        side_end_row = copy_block(
            source_packaging_start,
            source_summary_start - 2,
            side_start_column,
            side_end_column,
            packaging_title_row,
            link_computed=True,
        )

    lift_cost_row = packaging_usd_row + 3
    lift_markup_row = lift_cost_row + 1
    lift_settlement_row = lift_cost_row + 2
    lift_hkd_row = lift_cost_row + 3
    lift_fx_row = lift_cost_row + 4
    lift_usd_row = lift_cost_row + 5
    _template_cell(target, lift_cost_row, 3, "吊柜费", horizontal="left", border=None)
    _template_cell(target, lift_markup_row, 3, "×", horizontal="right", border=None)
    _template_cell(target, lift_settlement_row, 3, "÷", horizontal="right", border=Border(bottom=TEMPLATE_MEDIUM))
    _template_cell(target, lift_hkd_row, 3, "单价（HK$）：", bold=True, color=TEMPLATE_BLUE, horizontal="right", border=None)
    _template_cell(target, lift_fx_row, 3, "美金兑港币汇率：", color=TEMPLATE_BLUE, horizontal="right", border=None)
    _template_cell(target, lift_usd_row, 3, "吊柜费（USD）：", fill="FFFF00", bold=True, color=TEMPLATE_BLUE, horizontal="right", border=None)
    _template_cell(
        target,
        lift_markup_row,
        4,
        packaging_markup,
        number_format="0.000",
        border=None,
    )
    _template_cell(
        target,
        lift_settlement_row,
        4,
        "=1-$Q$6",
        number_format="0.0000",
        border=Border(bottom=TEMPLATE_MEDIUM),
    )
    for column, route in enumerate(route_rows, start=5):
        letter = get_column_letter(column)
        _template_cell(
            target,
            lift_cost_row,
            column,
            _float_value(route.get("lift_hkd")),
            number_format="0.000",
            border=None,
        )
        _template_cell(
            target,
            lift_markup_row,
            column,
            f"=$D${lift_markup_row}",
            number_format="0.000",
            border=None,
        )
        _template_cell(
            target,
            lift_settlement_row,
            column,
            f"=$D${lift_settlement_row}",
            number_format="0.0000",
            border=Border(bottom=TEMPLATE_MEDIUM),
        )
        _template_cell(
            target,
            lift_hkd_row,
            column,
            f"={letter}{lift_cost_row}*{letter}{lift_markup_row}/{letter}{lift_settlement_row}",
            bold=True,
            color=TEMPLATE_BLUE,
            number_format="0.00",
            border=None,
        )
        _template_cell(
            target,
            lift_fx_row,
            column,
            "=$R$4",
            color=TEMPLATE_BLUE,
            number_format="0.00",
            border=None,
        )
        _template_cell(
            target,
            lift_usd_row,
            column,
            f"={letter}{lift_hkd_row}/{letter}{lift_fx_row}",
            fill="FFFF00",
            bold=True,
            color=TEMPLATE_BLUE,
            number_format="0.000",
            border=None,
        )
    _apply_outline_border(
        target,
        lift_cost_row,
        lift_usd_row,
        3,
        quote_price_end_column,
        TEMPLATE_MEDIUM,
    )

    final_header_row = lift_usd_row + 3
    freight_total_row = final_header_row + 1
    combined_total_row = final_header_row + 3
    _template_cell(
        target,
        final_header_row,
        3,
        "",
        bold=True,
        wrap_text=False,
    )
    _template_cell(
        target,
        final_header_row,
        4,
        "出厂价",
        bold=True,
        wrap_text=False,
    )
    for index, route in enumerate(route_rows, start=5):
        _template_cell(
            target,
            final_header_row,
            index,
            _safe_text(route.get("name") or route.get("item"))
            or f"运输方案{index - 4}",
            bold=True,
            wrap_text=False,
        )
    for row_index, label in (
        (freight_total_row, "产品价（含运费）（USD）："),
        (combined_total_row, "产品价（含吊柜费）（USD）："),
    ):
        _template_cell(
            target,
            row_index,
            3,
            label,
            fill="FFFF00",
            bold=True,
            color=TEMPLATE_BLUE,
            horizontal="right",
            wrap_text=False,
            border=None,
        )
    component_usd_terms = "+".join(f"$D${row}" for row in component_usd_rows) or "0"
    for column in quote_price_columns:
        letter = get_column_letter(column)
        _template_cell(
            target,
            freight_total_row,
            column,
            f"={component_usd_terms}+{letter}{packaging_usd_row}",
            fill="FFFF00",
            bold=True,
            color=TEMPLATE_BLUE,
            number_format="0.00",
            wrap_text=False,
        )
        _template_cell(
            target,
            combined_total_row,
            column,
            (
                f"={letter}{freight_total_row}"
                if column == 4
                else f"={letter}{freight_total_row}+{letter}{lift_usd_row}"
            ),
            fill="FFFF00",
            bold=True,
            color=TEMPLATE_BLUE,
            number_format="0.00",
            wrap_text=False,
        )
    _apply_outline_border(
        target,
        final_header_row,
        combined_total_row,
        3,
        quote_price_end_column,
        TEMPLATE_MEDIUM,
    )

    selected_summary_column = quote_price_columns[-1]
    selected_summary_letter = get_column_letter(selected_summary_column)
    selected_route = route_rows[-1] if route_rows else {}
    summary_destination_row = max(combined_total_row, side_end_row) + 3
    summary_end_row = summary_destination_row - 1
    if source_summary_start <= source.max_row:
        summary_end_row = copy_block(
            source_summary_start,
            source.max_row,
            1,
            outer_right_column,
            summary_destination_row,
            link_computed=True,
        )

        first_header_row = summary_destination_row
        first_value_row = first_header_row + 1
        second_header_row = first_header_row + 3
        second_value_row = second_header_row + 1
        third_header_row = first_header_row + 6
        third_value_row = third_header_row + 1
        tax_header_row = first_header_row + 9
        tax_amount_row = tax_header_row + 1
        deduction_row = tax_header_row + 3

        def visible_sumif(
            criteria_column: str,
            criteria: str,
            *,
            amount_column: str = "D",
        ) -> str:
            return "=" + "+".join(
                f"SUMIF(${criteria_column}${start}:${criteria_column}${end},"
                f"{criteria},${amount_column}${start}:${amount_column}${end})"
                for start, end in visible_detail_ranges
            )

        def visible_sum_amounts() -> str:
            return "+".join(
                f"SUM($D${start}:$D${end})"
                for start, end in visible_detail_ranges
            ) or "0"

        for column in range(5, 15):
            target.cell(first_value_row, column).value = visible_sumif(
                "B", f"{get_column_letter(column)}{first_header_row}"
            )
        target.cell(first_value_row, 4).value = (
            f"={selected_summary_letter}{combined_total_row}*$R$4"
        )
        for column in range(6, 11):
            target.cell(second_value_row, column).value = visible_sumif(
                "B", f"{get_column_letter(column)}{second_header_row}"
            )
        target.cell(second_value_row, 11).value = _float_value(
            selected_route.get("freight_hkd")
        )
        target.cell(second_value_row, 12).value = _float_value(
            selected_route.get("lift_hkd")
        )
        target.cell(second_value_row, 13).value = f"=D{first_value_row}*$Q$6"
        additional_tax_formula = visible_sumif("C", '"附加税"')[1:]
        indonesia_formula = visible_sumif("C", '"印尼运费"')[1:]
        target.cell(second_value_row, 14).value = (
            f"=SUM(E{first_value_row}:N{first_value_row},"
            f"F{second_value_row}:M{second_value_row},G{third_value_row})"
            f"+({additional_tax_formula})+({indonesia_formula})"
        )
        target.cell(second_value_row, 4).value = (
            f"=IFERROR(D{first_value_row}/N{third_value_row},0)"
        )
        target.cell(second_value_row, 5).value = (
            f"=IFERROR(D{first_value_row}/P{deduction_row},0)"
        )

        abs_terms = [
            f'SUMIF($D${start}:$D${end},"*ABS*",$K${start}:$K${end})'
            for start, end in component_mold_ranges
        ]
        target.cell(second_value_row, 3).value = (
            "=" + "+".join(abs_terms) if abs_terms else "=0"
        )
        for column in range(5, 9):
            target.cell(third_value_row, column).value = visible_sumif(
                "B", f"{get_column_letter(column)}{third_header_row}"
            )
        target.cell(third_value_row, 3).value = (
            f"=IFERROR(C{second_value_row}/D{first_value_row},0)"
        )
        target.cell(third_value_row, 9).value = (
            f"=IFERROR((E{third_value_row}+F{third_value_row}+H{third_value_row})/"
            f"D{first_value_row},0)"
        )
        target.cell(third_value_row, 10).value = (
            f"=D{first_value_row}-N{second_value_row}"
        )
        target.cell(third_value_row, 11).value = (
            f"=IFERROR(J{third_value_row}/D{first_value_row},0)"
        )
        target.cell(third_value_row, 12).value = (
            f"=D{first_value_row}-N{third_value_row}"
        )
        target.cell(third_value_row, 13).value = (
            f"=IFERROR(L{third_value_row}/D{first_value_row},0)"
        )
        target.cell(third_value_row, 14).value = (
            f"=N{second_value_row}+E{third_value_row}+F{third_value_row}+H{third_value_row}"
        )
        target.cell(third_value_row, 15).value = (
            f"={visible_sum_amounts()}+K{second_value_row}+L{second_value_row}"
            f"+M{second_value_row}"
        )

        tax_13_formula = visible_sumif("A", '"¥13%"')
        tax_1_formula = visible_sumif("A", '"¥1%"')[1:]
        target.cell(tax_amount_row, 4).value = tax_13_formula
        target.cell(tax_amount_row, 3).value = (
            f"=D{tax_amount_row}+({tax_1_formula})"
        )
        t4 = _rr2_tax_rows(rr2_cost_summary)
        tax_keys = (
            "tax1",
            "slush3",
            "sewhair13",
            "sewcloth13",
            "suction6",
            "freight9",
            "tax13b",
            "carton",
            "labor13",
        )
        tax_amount_references = {
            "tax1": f"H{second_value_row}",
            "slush3": f"G{first_value_row}",
            "sewhair13": f"H{first_value_row}",
            "suction6": f"M{first_value_row}",
            "freight9": f"K{second_value_row}",
            "tax13b": f"D{tax_amount_row}",
            "carton": f"J{second_value_row}",
            "labor13": f"SUM(E{third_value_row}:F{third_value_row},H{third_value_row})",
        }
        for column, key in enumerate(tax_keys, start=6):
            rate_percent = t4.get(key, {}).get("rate_percent")
            if key == "carton" or (
                key == "sewcloth13" and rate_percent in (None, "")
            ):
                target.cell(deduction_row, column).value = ""
                continue
            amount_reference = (
                format(_float_value(t4.get(key, {}).get("amount_hkd")), ".10g")
                if key == "sewcloth13"
                else tax_amount_references[key]
            )
            target.cell(deduction_row, column).value = (
                f"={amount_reference}*{get_column_letter(column)}{tax_amount_row}"
            )
        target.cell(deduction_row, 15).value = (
            f"=SUM(F{deduction_row}:N{deduction_row})"
        )
        target.cell(deduction_row, 16).value = (
            f"=N{third_value_row}-O{deduction_row}"
        )

    first_component_end = component_usd_rows[0] if component_usd_rows else packaging_title_row - 1
    _add_product_image_to_region(
        target,
        attachments,
        min_row=8,
        max_row=max(8, first_component_end),
        min_column=side_start_column,
        max_column=side_end_column,
    )
    target.freeze_panes = "A8"
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.sheet_properties.pageSetUpPr.fitToPage = True
    target.page_setup.orientation = "landscape"
    target.page_setup.paperSize = target.PAPERSIZE_A3
    target.page_setup.fitToWidth = 1
    target.page_setup.fitToHeight = 0
    target.page_margins = copy(source.page_margins)
    target.print_area = f"A1:{get_column_letter(outer_right_column)}{summary_end_row}"
    target.sheet_properties.outlinePr.summaryBelow = True
    _apply_outline_border(
        target,
        1,
        summary_end_row,
        1,
        outer_right_column,
        TEMPLATE_MEDIUM,
    )
    workbook.active = 0


def _walk_components(rows: list[dict[str, Any]], parent: str = ""):
    for row in rows:
        if not isinstance(row, dict):
            continue
        yield parent, row
        children = row.get("children", [])
        if isinstance(children, list):
            yield from _walk_components(children, str(row.get("item", "")))


def _build_electronic_sheet(
    workbook: Workbook,
    section: InternalQuoteSection | None,
    reference_snapshot: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet("电子明细")
    _style_title(sheet, "电子报价明细", 10)
    _header_row(sheet, 3, ("父项", "零件名称", "规格", "用量", "单价RMB", "单价HKD", "金额HKD", "税点%", "备注", "来源"))
    row_index = 4
    payload = _json_object(section.payload_json) if section else {}
    fx_value = _number(reference_snapshot.get("fx", {}).get("rmb_hkd"))
    fx = fx_value if isinstance(fx_value, float) and fx_value > 0 else 0.85
    components = payload.get("components", [])
    for parent, row in _walk_components(components if isinstance(components, list) else []):
        quantity = _number(row.get("quantity"))
        unit_price_rmb = _number(row.get("unit_price_rmb"))
        unit_price_hkd = _number(row.get("unit_price_hkd"))
        if not isinstance(unit_price_rmb, float) and isinstance(unit_price_hkd, float):
            unit_price_rmb = unit_price_hkd * fx
        if not isinstance(unit_price_hkd, float) and isinstance(unit_price_rmb, float):
            unit_price_hkd = unit_price_rmb / fx
        amount_hkd = quantity * unit_price_hkd if isinstance(quantity, float) and isinstance(unit_price_hkd, float) else ""
        _body_row(
            sheet,
            row_index,
            (
                _safe_text(parent),
                _safe_text(row.get("item", "")),
                _safe_text(row.get("specification", "")),
                quantity,
                unit_price_rmb,
                unit_price_hkd,
                amount_hkd,
                _number(row.get("tax_rate_percent")),
                _safe_text(row.get("remark", row.get("note", ""))),
                _safe_text(f"{row.get('source_currency', '')} row {row.get('source_row', '')}"),
            ),
            amount_columns={5, 6, 7, 8},
        )
        row_index += 1
    calculation = _json_object(section.calculation_json) if section else {}
    totals = calculation.get("totals", {}) if isinstance(calculation.get("totals", {}), dict) else {}
    row_index += 1
    _header_row(sheet, row_index, ("电子成本汇总", "RMB", "HKD", "公式口径"))
    row_index += 1
    for label, rmb_key, hkd_key, formula in (
        ("零件成本", "component_rmb", "component_hkd", "用量 × 单价"),
        ("成本合计（不含税）", "pre_tax_rmb", "pre_tax_hkd", "零件 + 邦定/贴片/人工/测试/包装"),
        ("含利润价", "with_profit_rmb", "with_profit_hkd", "不含税成本 × (1 + 利润率)"),
        ("抵税差额", "tax_credit_difference_rmb", "tax_credit_difference_hkd", "含利润价 × 13% - 零件进项抵扣"),
        ("应交税负", "tax_payable_rmb", "tax_payable_hkd", "抵税差额 × 10%"),
        ("含税报价", "total_rmb", "total_hkd", "含利润价 + 抵税差额 + 应交税负"),
    ):
        _body_row(sheet, row_index, (label, _number(totals.get(rmb_key)), _number(totals.get(hkd_key)), formula), amount_columns={2, 3})
        row_index += 1
    _finish_sheet(sheet, (20, 28, 28, 12, 15, 15, 15, 12, 28, 20))


def _build_sewing_sheet(
    workbook: Workbook,
    section: InternalQuoteSection | None,
    reference_snapshot: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet("车缝明细")
    _style_title(sheet, "车缝报价明细", 12)
    _header_row(
        sheet,
        3,
        (
            "物料名称",
            "裁片部位",
            "供应商",
            "布料MOQ/Y",
            "低于MOQ/每色费用 RMB",
            "用量/码",
            "单价 RMB",
            "汇率",
            "成本 HKD",
            "码点",
            "价钱 HKD",
            "备注",
        ),
    )
    row_index = 4
    payload = _json_object(section.payload_json) if section else {}
    fx = _float_value(_dict_value(reference_snapshot.get("fx", {})).get("rmb_hkd"), 0.85)
    if fx <= 0:
        fx = 0.85
    for group in payload.get("groups", []) if isinstance(payload.get("groups", []), list) else []:
        if not isinstance(group, dict):
            continue
        group_name = _safe_text(group.get("name")) or "车缝产品组"
        group_category = "车发" if str(group.get("category")) == "hair" else "车衣"
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=12)
        group_cell = sheet.cell(row_index, 1)
        group_cell.value = f"{group_name} · {group_category}"
        group_cell.fill = PatternFill("solid", fgColor="F8CBAD")
        group_cell.font = Font(name="宋体", size=10, bold=True, color="7F6000")
        group_cell.alignment = Alignment(horizontal="left", vertical="center")
        group_cell.border = Border(
            left=TEMPLATE_THIN,
            right=TEMPLATE_THIN,
            top=TEMPLATE_THIN,
            bottom=TEMPLATE_THIN,
        )
        sheet.row_dimensions[row_index].height = 20
        row_index += 1
        for row in group.get("materials", []) if isinstance(group.get("materials", []), list) else []:
            if not isinstance(row, dict):
                continue
            usage = _number(row.get("usage"))
            unit_price_rmb = _number(row.get("unit_price_rmb"))
            exchange_rate = _number(row.get("exchange_rate"))
            if not isinstance(exchange_rate, float) or exchange_rate <= 0:
                exchange_rate = fx
            markup = _number(row.get("markup")) or 1
            extra_evidence = []
            if row.get("craft"):
                extra_evidence.append(f"工艺：{_safe_text(row.get('craft'))}")
            if _float_value(row.get("pieces")) > 0:
                extra_evidence.append(f"裁片数：{_compact_number_label(row.get('pieces'))}")
            if row.get("source_row") not in (None, ""):
                extra_evidence.append(f"来源行：{_safe_text(row.get('source_row'))}")
            remark = _safe_text(row.get("remark") or row.get("note") or "")
            if extra_evidence:
                remark = "；".join([remark, *extra_evidence] if remark else extra_evidence)
            _body_row(
                sheet,
                row_index,
                (
                    _safe_text(row.get("item", "")),
                    _safe_text(row.get("part", "")),
                    _safe_text(row.get("supplier", "")),
                    _number(row.get("fabric_moq_y")),
                    _number(row.get("below_moq_fee_rmb")),
                    usage,
                    unit_price_rmb,
                    exchange_rate,
                    "",
                    markup,
                    "",
                    remark,
                ),
                amount_columns={4, 5, 6, 7, 8, 9, 10, 11},
            )
            sheet.cell(row_index, 9).value = f"=F{row_index}*G{row_index}/H{row_index}"
            sheet.cell(row_index, 9).number_format = "#,##0.0000"
            sheet.cell(row_index, 11).value = f"=I{row_index}*J{row_index}"
            sheet.cell(row_index, 11).number_format = "#,##0.0000"
            row_index += 1
    if row_index > 4:
        _body_row(sheet, row_index, ("", "", "", "", "", "", "", "", "", "合计", "", ""))
        sheet.cell(row_index, 11).value = f"=SUM(K4:K{row_index - 1})"
        sheet.cell(row_index, 11).number_format = "#,##0.0000"
        sheet.cell(row_index, 10).fill = PatternFill("solid", fgColor="FFF200")
        sheet.cell(row_index, 11).fill = PatternFill("solid", fgColor="FFF200")
        sheet.cell(row_index, 10).font = Font(name="宋体", size=10, bold=True)
        sheet.cell(row_index, 11).font = Font(name="宋体", size=10, bold=True)
    _finish_sheet(sheet, (34, 18, 18, 15, 24, 14, 14, 12, 15, 12, 15, 36))


def _build_hair_sheet(workbook: Workbook, section: InternalQuoteSection | None) -> None:
    sheet = workbook.create_sheet("车发明细")
    _style_title(sheet, "车发报价明细", 8)
    _header_row(sheet, 3, ("#", "名称", "工艺", "重量(g)", "单价(HKD)", "单位", "备注", "金额(HKD)"))
    row_index = 4
    payload = _json_object(section.payload_json) if section else {}
    calculation = _json_object(section.calculation_json) if section else {}
    breakdown = calculation.get("line_breakdown", [])
    amounts = [
        _number(row.get("amount_hkd"))
        for row in breakdown
        if isinstance(row, dict) and row.get("kind") == "hair"
    ] if isinstance(breakdown, list) else []
    for item_index, row in enumerate(
        payload.get("lines", []) if isinstance(payload.get("lines", []), list) else [],
        start=1,
    ):
        if not isinstance(row, dict):
            continue
        amount = amounts[item_index - 1] if item_index - 1 < len(amounts) else _number(row.get("unit_price_hkd"))
        _body_row(
            sheet,
            row_index,
            (
                item_index,
                _safe_text(row.get("name", row.get("item", ""))),
                _safe_text(row.get("craft", row.get("process", ""))),
                _number(row.get("weight_g")),
                _number(row.get("unit_price_hkd")),
                _safe_text(row.get("unit", "")),
                _safe_text(row.get("remark", row.get("note", ""))),
                amount,
            ),
            amount_columns={4, 5, 8},
        )
        row_index += 1
    _finish_sheet(sheet, (8, 28, 24, 14, 16, 14, 32, 16))


def _build_assembly_sheet(workbook: Workbook, section: InternalQuoteSection | None) -> None:
    sheet = workbook.create_sheet("装配明细")
    _style_title(sheet, "装配排拉工序明细", 8)
    _header_row(sheet, 3, ("产品组", "分类", "工序", "人数", "小组数", "生产量", "备注", "来源行"))
    row_index = 4
    payload = _json_object(section.payload_json) if section else {}
    for group in payload.get("groups", []) if isinstance(payload.get("groups", []), list) else []:
        if not isinstance(group, dict):
            continue
        for process in group.get("processes", []) if isinstance(group.get("processes", []), list) else []:
            if not isinstance(process, dict):
                continue
            _body_row(
                sheet,
                row_index,
                (
                    _safe_text(group.get("name", "")),
                    _safe_text(group.get("category", "")),
                    _safe_text(process.get("name", "")),
                    _number(process.get("persons")),
                    _number(process.get("teams")),
                    _number(process.get("production_qty")),
                    _safe_text(process.get("note", "")),
                    process.get("source_row", ""),
                ),
            )
            row_index += 1
    _finish_sheet(sheet, (24, 14, 30, 12, 12, 15, 32, 12))


def _build_structured_data_sheet(
    workbook: Workbook,
    quote: InternalQuote,
    sections: list[InternalQuoteSection],
    reference_snapshot: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet("结构化数据")
    _style_title(sheet, "P4 客价转换结构化数据", 12)
    _body_row(
        sheet,
        2,
        (
            "结构版本",
            STRUCTURED_DATA_SCHEMA_VERSION,
            "说明",
            "按记录类型、分段和分片序号重组 JSON；禁止从展示明细反推原始参数",
        ),
    )
    _header_row(
        sheet,
        3,
        (
            "记录类型",
            "分段代码",
            "分段名称",
            "状态",
            "revision",
            "计算状态",
            "依赖状态",
            "计算hash",
            "分片序号",
            "分片总数",
            "JSON分片",
            "是否参与",
        ),
    )
    row_index = 4
    snapshot_chunks = _json_chunks(reference_snapshot)
    for chunk_index, chunk in enumerate(snapshot_chunks, start=1):
        _body_row(
            sheet,
            row_index,
            (
                "reference_snapshot",
                "quote",
                "报价参考快照",
                "frozen",
                quote.header_revision,
                "valid",
                "current",
                quote.reference_snapshot_id,
                chunk_index,
                len(snapshot_chunks),
                chunk,
                "是",
            ),
            text_columns={11},
        )
        row_index += 1
    by_code = {section.department: section for section in sections}
    for code in SECTION_ORDER:
        section = by_code.get(code)
        if section is None:
            continue
        records = (
            ("payload", _json_object(section.payload_json)),
            ("calculation", _json_object(section.calculation_json)),
        )
        for record_type, value in records:
            chunks = _json_chunks(value)
            for chunk_index, chunk in enumerate(chunks, start=1):
                _body_row(
                    sheet,
                    row_index,
                    (
                        record_type,
                        section.department,
                        _safe_text(section.department_name),
                        section.status,
                        section.revision,
                        section.calculation_status,
                        section.dependency_status,
                        section.calculation_hash,
                        chunk_index,
                        len(chunks),
                        chunk,
                        "是" if section.is_required else "否",
                    ),
                    text_columns={11},
                )
                row_index += 1
    _finish_sheet(sheet, (16, 16, 20, 18, 10, 16, 16, 66, 10, 10, 76, 12))


def _build_approval_sheet(
    workbook: Workbook,
    quote: InternalQuote,
    sections: list[InternalQuoteSection],
    manifest: dict[str, Any],
) -> None:
    sheet = workbook.create_sheet("审批与版本")
    _style_title(sheet, "审批、公式与版本清单", 8)
    is_final_release = manifest.get("release_stage") == "p4_final_approved"
    template_version = str(manifest.get("template_version") or P3_TEMPLATE_VERSION)
    release_label = "P4 最终业务放行" if is_final_release else "P3 分段审批后内部成本快照"
    boundary_label = (
        "最终业务放行完成，可交接客价转换台"
        if is_final_release
        else "最终业务放行与客价交接在 P4 实施"
    )
    summary = (
        ("模板版本", template_version, "公式版本", quote.formula_version),
        ("参考快照", quote.reference_snapshot_id, "报价头revision", quote.header_revision),
        ("导出阶段", release_label, "清单SHA-256", manifest.get("manifest_sha256", "")),
        ("边界说明", boundary_label, "", ""),
        (
            "最终提交人",
            manifest.get("final_submitted_by_name", ""),
            "最终放行人/时间",
            f"{manifest.get('final_reviewed_by_name', '')} {manifest.get('final_reviewed_at', '')}".strip(),
        ),
    )
    for row_index, values in enumerate(summary, start=2):
        _body_row(sheet, row_index, (_safe_text(value) for value in values))
    _header_row(sheet, 7, ("分段", "是否必需", "状态", "revision", "计算状态", "依赖状态", "审核人", "审核时间"))
    row_index = 8
    by_code = {section.department: section for section in sections}
    for code in SECTION_ORDER:
        section = by_code.get(code)
        if section is None:
            continue
        _body_row(
            sheet,
            row_index,
            (
                _safe_text(section.department_name),
                "是" if section.is_required else "否",
                section.status,
                section.revision,
                section.calculation_status,
                section.dependency_status,
                _safe_text(section.reviewed_by),
                section.reviewed_at,
            ),
        )
        row_index += 1
    row_index += 1
    _header_row(sheet, row_index, ("清单字段", "值"))
    row_index += 1
    for key, value in manifest.items():
        _body_row(sheet, row_index, (_safe_text(key), _safe_text(_json_text(value) if isinstance(value, (dict, list)) else value)))
        row_index += 1
    _finish_sheet(sheet, (24, 24, 20, 16, 18, 18, 20, 22))


def _unique_sheet_title(raw_title: str, used_titles: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\\x00-\x1f]", " ", raw_title).strip().strip("'")
    cleaned = re.sub(r"\s+", " ", cleaned) or "上传表格"
    base = cleaned[:31]
    candidate = base
    suffix = 2
    while candidate.casefold() in used_titles:
        marker = f"-{suffix}"
        candidate = f"{base[:31 - len(marker)]}{marker}"
        suffix += 1
    used_titles.add(candidate.casefold())
    return candidate


def _rewrite_attachment_formula(
    formula: object,
    sheet_names: dict[str, str],
) -> object:
    if not isinstance(formula, str) or not formula.startswith("="):
        return formula
    rendered = formula
    for old_name, new_name in sheet_names.items():
        if old_name == new_name:
            continue
        old_quoted = f"{quote_sheetname(old_name)}!"
        new_quoted = f"{quote_sheetname(new_name)}!"
        rendered = rendered.replace(old_quoted, new_quoted)
        rendered = rendered.replace(f"{old_name}!", new_quoted)
    return rendered


def _copy_attachment_worksheet(
    source,
    target,
    sheet_names: dict[str, str],
) -> None:
    for row in source.iter_rows():
        for source_cell in row:
            target_cell = target.cell(source_cell.row, source_cell.column)
            target_cell.value = _rewrite_attachment_formula(
                source_cell.value,
                sheet_names,
            )
            if source_cell.has_style:
                # Style-array indexes are workbook-local.  Copy the concrete
                # objects so openpyxl registers them in the destination book.
                target_cell.font = copy(source_cell.font)
                target_cell.fill = copy(source_cell.fill)
                target_cell.border = copy(source_cell.border)
                target_cell.alignment = copy(source_cell.alignment)
                target_cell.protection = copy(source_cell.protection)
                target_cell.number_format = source_cell.number_format
            if source_cell.hyperlink is not None:
                target_cell._hyperlink = copy(source_cell.hyperlink)
            if source_cell.comment is not None:
                target_cell.comment = copy(source_cell.comment)

    for merged_range in source.merged_cells.ranges:
        target.merge_cells(str(merged_range))
    for index, dimension in source.row_dimensions.items():
        target_dimension = target.row_dimensions[index]
        target_dimension.height = dimension.height
        target_dimension.hidden = dimension.hidden
        target_dimension.outlineLevel = dimension.outlineLevel
        target_dimension.collapsed = dimension.collapsed
    for index, dimension in source.column_dimensions.items():
        target_dimension = target.column_dimensions[index]
        target_dimension.width = dimension.width
        target_dimension.hidden = dimension.hidden
        target_dimension.bestFit = dimension.bestFit
        target_dimension.outlineLevel = dimension.outlineLevel
        target_dimension.collapsed = dimension.collapsed

    target.freeze_panes = source.freeze_panes
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.sheet_view.zoomScale = source.sheet_view.zoomScale
    target.sheet_format.defaultRowHeight = source.sheet_format.defaultRowHeight
    target.sheet_format.defaultColWidth = source.sheet_format.defaultColWidth
    target.page_margins = copy(source.page_margins)
    target.page_setup = copy(source.page_setup)
    target.print_options = copy(source.print_options)
    target.sheet_properties = copy(source.sheet_properties)
    target.auto_filter.ref = source.auto_filter.ref
    target.data_validations = copy(source.data_validations)
    target.conditional_formatting = copy(source.conditional_formatting)

    for image in getattr(source, "_images", []):
        try:
            copied_image = copy(image)
            copied_image.anchor = copy(image.anchor)
            target.add_image(copied_image)
        except Exception:
            # Cell values and styles remain authoritative even when a source
            # drawing uses a format that openpyxl cannot clone across books.
            continue
    for chart in getattr(source, "_charts", []):
        try:
            copied_chart = copy(chart)
            copied_chart.anchor = copy(chart.anchor)
            target.add_chart(copied_chart)
        except Exception:
            continue


def _append_spreadsheet_attachment_sheets(
    workbook: Workbook,
    attachments: list[InternalQuoteAttachment],
) -> list[str]:
    appended_titles: list[str] = []
    used_titles = {sheet.title.casefold() for sheet in workbook.worksheets}
    supported_suffixes = {".xlsx", ".xlsm"}
    seen_hashes: set[str] = set()
    for attachment in attachments:
        suffix = Path(attachment.file_name).suffix.lower()
        if suffix not in supported_suffixes or attachment.sha256 in seen_hashes:
            continue
        seen_hashes.add(attachment.sha256)
        try:
            source_workbook = load_workbook(
                BytesIO(attachment.content),
                data_only=False,
                read_only=False,
                keep_links=False,
            )
        except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError):
            title = _unique_sheet_title(
                f"{Path(attachment.file_name).stem}-附件异常",
                used_titles,
            )
            sheet = workbook.create_sheet(title)
            _style_title(sheet, "上传表格未能嵌入", 4)
            _body_row(
                sheet,
                3,
                (
                    "文件名",
                    _safe_text(attachment.file_name),
                    "SHA-256",
                    _safe_text(attachment.sha256),
                ),
            )
            _body_row(
                sheet,
                4,
                ("说明", "该文件仍保留在系统附件中，请重新另存为有效 xlsx 后上传。"),
            )
            _finish_sheet(sheet, (18, 56, 18, 66))
            appended_titles.append(title)
            continue

        try:
            source_sheets = list(source_workbook.worksheets)
            stem = Path(attachment.file_name).stem
            sheet_names: dict[str, str] = {}
            for source_sheet in source_sheets:
                raw_title = (
                    stem
                    if len(source_sheets) == 1
                    else f"{stem}-{source_sheet.title}"
                )
                sheet_names[source_sheet.title] = _unique_sheet_title(
                    raw_title,
                    used_titles,
                )
            for source_sheet in source_sheets:
                target = workbook.create_sheet(sheet_names[source_sheet.title])
                _copy_attachment_worksheet(source_sheet, target, sheet_names)
                target.sheet_state = "visible"
                appended_titles.append(target.title)
        finally:
            source_workbook.close()
    return appended_titles


def build_internal_quote_workbook(
    quote: InternalQuote,
    sections: list[InternalQuoteSection],
    manifest: dict[str, Any],
    reference_snapshot: dict[str, Any] | None = None,
    rr2_cost_summary: dict[str, Any] | None = None,
    cost_context: dict[str, Any] | None = None,
    attachments: list[InternalQuoteAttachment] | None = None,
) -> bytes:
    workbook = Workbook()
    workbook.properties.creator = "Royal Regent Nexus"
    workbook.properties.title = f"{quote.quote_no} 内部报价"
    _build_summary_sheet(
        workbook,
        quote,
        sections,
        reference_snapshot or {},
        rr2_cost_summary or {},
        cost_context or {},
        attachments or [],
    )
    by_code = {section.department: section for section in sections}
    _build_electronic_sheet(workbook, by_code.get("electronic"), reference_snapshot or {})
    _build_sewing_sheet(workbook, by_code.get("sewing"), reference_snapshot or {})
    _build_hair_sheet(workbook, by_code.get("hair"))
    _build_assembly_sheet(workbook, by_code.get("assembly"))
    if manifest.get("release_stage") == "p4_final_approved":
        _build_structured_data_sheet(workbook, quote, sections, reference_snapshot or {})
    _build_approval_sheet(workbook, quote, sections, manifest)
    for technical_sheet in workbook.worksheets[1:]:
        technical_sheet.sheet_state = "veryHidden"
    _append_spreadsheet_attachment_sheets(workbook, attachments or [])
    workbook.active = 0
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
