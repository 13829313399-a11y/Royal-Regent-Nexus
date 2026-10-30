"""Cumulative OOXML writer. Preserve untouched ZIP parts and original history.

Only owned rows/sheets are regenerated. Historical bindings keep their original
row XML so moving an inspection never destroys another historical record.
"""
from collections import defaultdict
from copy import deepcopy
from datetime import date, datetime, timedelta
from io import BytesIO
import json
import math
import posixpath
import re
import unicodedata
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E

from fastapi import HTTPException
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.datetime import from_excel, to_excel

S='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P='http://schemas.openxmlformats.org/package/2006/relationships'
C='http://schemas.openxmlformats.org/package/2006/content-types'
D='http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'
A='http://schemas.openxmlformats.org/drawingml/2006/main'
for prefix,ns in [('r',R),('xdr',D),('a',A)]:E.register_namespace(prefix,ns)
def tag(name):return '{'+S+'}'+name
def xml(node):return E.tostring(node,encoding='utf-8',xml_declaration=True)
def text(value):return '' if value is None else str(value)
def normalized(value):
    if isinstance(value,float) and value.is_integer():value=int(value)
    return ' '.join(unicodedata.normalize('NFKC',text(value)).split()).casefold()

CLIENTS=['ZURU','Dickie','EDU','彩星','360','SHUSHUPAPA','MMI','银辉','MAXX','BUZZBEE','SEASONS']
TARGETS={normalized(n):n+('' if n in ('MMI','银辉','MAXX') else '行验货') for n in CLIENTS}
TARGETS['buzz bee']='BUZZBEE行验货'
SECOND={'BUZZBEE行验货','SEASONS行验货'}
PROTECTED={'新明自检','迪奇自检','新明退货','退货未报','Dickie行验货 (2)'}
HEAD1=['客户','合同','出口国','货号','产品名称','数量','装箱','箱数','结果','验货问题','第一责任人','第二责任人','备注']
HEAD2=['Contract No.','客名/國家','產品編號','產品名称','數量','装箱','箱数','结果','工厂','报告','问题','责任人','备注']
FORMAT_VERSION=5
SUMMARY=['日期','客户','PO#','出口国','货号','产品名称','数量','装箱','箱数','生产','机构','客PO#','完成情况','备注']

def target(item):
    customer=text(item.get('customer')).strip()
    if not customer:raise HTTPException(422,f"记录 {item['id']} 缺少客户")
    result=TARGETS.get(normalized(customer))
    if result:return result
    if normalized(customer) in {normalized(x) for x in PROTECTED}|{'新明','迪奇'}:
        raise HTTPException(422,f"记录 {item['id']} 的业务类型未明确：{customer}")
    result=customer+'行验货'
    if len(result)>31 or re.search(r'[\\/*?:\[\]]',result) or result in PROTECTED:
        raise HTTPException(422,f"记录 {item['id']} 无法唯一确定客户工作表：{customer}")
    return result

def numbers(item):
    q,u,c=[item.get(k,'') for k in ('quantity','unitsPerCarton','cartonCount')]
    def number(x):
        if isinstance(x,bool):return None
        if not re.fullmatch(r'\d+(?:\.\d+)?',text(x).strip()):return None
        v=float(x)
        return v if math.isfinite(v) else None
    if c in ('',None) and number(q) is not None and number(u) not in (None,0):c=number(q)/number(u)
    if all(number(v) is not None for v in (q,u,c)) and number(u)>0 and abs(number(q)-number(u)*number(c))>1e-6:
        raise HTTPException(422,f"记录 {item['id']} 的数量、装箱、箱数不一致，未自动改值")
    return [number(v) if number(v) is not None else text(v) for v in (q,u,c)]

def client_values(item,second=False):
    q,u,c=numbers(item)
    problem=item.get('inspectionProblem',item.get('remark',''))
    other=item.get('otherRemark','')
    if second:return [item.get('poNo'),item.get('country'),item.get('productNo'),item.get('productName'),q,u,c,item.get('result'),item.get('inspectionFactory',''),item.get('reportStatus',''),problem,item.get('responsible',''),other]
    return [item.get('customer'),item.get('poNo'),item.get('country'),item.get('productNo'),item.get('productName'),q,u,c,item.get('result'),problem,item.get('firstResponsible',''),item.get('secondResponsible',''),other]

def displayed_date(item):
    if item.get('actualDate'):return date.fromisoformat(item['actualDate'])
    start,end=item.get('dateFrom',''),item.get('dateTo','')
    return start+' 至 '+end if end and end!=start else date.fromisoformat(start)

def record_order(pair):
    identity,item=pair
    return (item['archiveDate'],normalized(item.get('customer')),normalized(item.get('poNo')),normalized(item.get('productNo')),identity)

def data_styles(p,profile,second=False):
    wrap={1,2,3,4,9,10,11,12,13} if second else {1,2,3,4,5,10,11,12,13}
    return {c:p.style(v,c in wrap) for c,v in profile['data'].items()}

class Package:
    def __init__(self,content):
        with ZipFile(BytesIO(content)) as z:self.parts={n:z.read(n) for n in z.namelist()}
        self.book=E.fromstring(self.parts['xl/workbook.xml']);self.sheets=self.book.find(tag('sheets'))
        self.original_order=[s.get('name') for s in self.sheets]
        self.rels=E.fromstring(self.parts['xl/_rels/workbook.xml.rels'])
        self.types=E.fromstring(self.parts['[Content_Types].xml'])
        self.styles=E.fromstring(self.parts['xl/styles.xml'])
        self.shared=[]
        if 'xl/sharedStrings.xml' in self.parts:
            self.shared=[''.join(e.itertext()) for e in E.fromstring(self.parts['xl/sharedStrings.xml'])]
        self.loaded={};self.wrap_styles={}
        props=self.book.find(tag('workbookPr'));self.epoch=datetime(1904,1,1) if props is not None and props.get('date1904') in ('1','true') else datetime(1899,12,30)
    def names(self):return [s.get('name') for s in self.sheets]
    def path(self,name):
        sheet=next(s for s in self.sheets if s.get('name')==name)
        rel=next(r for r in self.rels if r.get('Id')==sheet.get('{'+R+'}id'))
        p=rel.get('Target');return p.lstrip('/') if p.startswith('/') else posixpath.normpath('xl/'+p)
    def sheet(self,name):
        if name not in self.loaded:self.loaded[name]=E.fromstring(self.parts[self.path(name)])
        return self.loaded[name]
    def rows(self,name):return self.sheet(name).find(tag('sheetData'))
    def values(self,row):
        values={}
        for c in row.findall(tag('c')):
            v=c.find(tag('v'));value=v.text if v is not None else ''
            if c.get('t')=='s':value=self.shared[int(value)] if value else ''
            elif c.get('t')=='inlineStr':value=''.join(c.find(tag('is')).itertext()) if c.find(tag('is')) is not None else ''
            elif c.get('t') not in ('str','e','b') and value:
                try:value=float(value)
                except ValueError:pass
            values[column_index_from_string(re.sub(r'\d','',c.get('r')))]=value
        return values
    def add(self,name,template=None):
        if name in self.names():raise HTTPException(422,'新增工作表名称冲突：'+name)
        sid=max(int(s.get('sheetId')) for s in self.sheets)+1
        rid='qc'+str(sid)
        while any(r.get('Id')==rid for r in self.rels):rid+='x'
        path='xl/worksheets/qc'+str(sid)+'.xml'
        E.SubElement(self.sheets,tag('sheet'),{'name':name,'sheetId':str(sid),'{'+R+'}id':rid})
        E.SubElement(self.rels,'{'+P+'}Relationship',{'Id':rid,'Type':R+'/worksheet','Target':path[3:]})
        E.SubElement(self.types,'{'+C+'}Override',{'PartName':'/'+path,'ContentType':'application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml'})
        root=E.Element(tag('worksheet'))
        if template:
            for n in ('sheetFormatPr','cols'):
                e=self.sheet(template).find(tag(n))
                if e is not None:root.append(deepcopy(e))
        E.SubElement(root,tag('sheetData'))
        if template:
            for n in ('printOptions','pageMargins','pageSetup','headerFooter'):
                e=self.sheet(template).find(tag(n))
                if e is not None:root.append(deepcopy(e))
        self.loaded[name]=root
        return root
    def remove(self,name):
        defined=self.book.find(tag('definedNames'))
        if defined is not None and name in self.original_order:
            for node in list(defined):
                if node.get('localSheetId')==str(self.original_order.index(name)):defined.remove(node)
        path=self.path(name);sheet=next(s for s in self.sheets if s.get('name')==name)
        rid=sheet.get('{'+R+'}id');self.sheets.remove(sheet)
        for r in list(self.rels):
            if r.get('Id')==rid:self.rels.remove(r)
        for t in list(self.types):
            if t.get('PartName')=='/'+path:self.types.remove(t)
        self.parts.pop(path,None);self.loaded.pop(name,None)
        self.parts.pop(posixpath.dirname(path)+'/_rels/'+posixpath.basename(path)+'.rels',None)
    def style(self,style,wrap=True,num_format=None):
        if not wrap and num_format is None:return str(style)
        style=int(style or 0);key=(style,wrap,num_format)
        if key not in self.wrap_styles:
            xfs=self.styles.find(tag('cellXfs'));xf=deepcopy(xfs[style])
            if wrap:
                alignment=xf.find(tag('alignment'))
                if alignment is None:alignment=E.SubElement(xf,tag('alignment'))
                alignment.set('wrapText','1');alignment.set('shrinkToFit','0')
                xf.set('applyAlignment','1')
            if num_format is not None:xf.set('numFmtId',str(num_format));xf.set('applyNumberFormat','1')
            # Preserve font, borders and both alignment directions; reuse equivalent styles.
            found=next((str(i) for i,old in enumerate(xfs) if E.tostring(old)==E.tostring(xf)),None)
            if found is None:
                xfs.append(xf);xfs.set('count',str(len(xfs)));found=str(len(xfs)-1)
            self.wrap_styles[key]=found
        return self.wrap_styles[key]
    def date_style(self,style,date_profile,value):
        source=self.styles.find(tag('cellXfs'))[int(date_profile.get(1,'0'))]
        fmt=source.get('numFmtId','14')
        return self.style(style,True,(fmt if fmt!='0' else '14') if isinstance(value,date) else '49')
    def font_size(self,style):
        xf=self.styles.find(tag('cellXfs'))[int(style or 0)]
        font=self.styles.find(tag('fonts'))[int(xf.get('fontId','0'))]
        size=font.find(tag('sz'))
        return float(size.get('val','11')) if size is not None else 11
    def cell(self,row,col,value,style='0'):
        if len(text(value))>32767 or any(ord(c)<32 and c not in '\r\n\t' for c in text(value)):
            raise HTTPException(422,'单元格文字超长或含不可显示字符，未截断原文')
        c=E.SubElement(row,tag('c'),{'r':get_column_letter(col)+row.get('r'),'s':str(style)})
        if isinstance(value,date):value=to_excel(value,self.epoch)
        if isinstance(value,(int,float)) and math.isfinite(value):E.SubElement(c,tag('v')).text=str(value)
        else:
            c.set('t','inlineStr');t=E.SubElement(E.SubElement(c,tag('is')),tag('t'))
            t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=text(value)
        return c
    def second(self,name):
        if name in getattr(self,'layouts',{}):return self.layouts[name]==2
        for row in self.rows(name):
            first=normalized(self.values(row).get(1))
            if first==normalized('Contract No.'):return True
            if first==normalized('客户'):return False
        return name in SECOND
    def client_names(self):
        return [name for name in self.names() if name not in PROTECTED and any(normalized(self.values(row).get(1)) in (normalized('客户'),normalized('Contract No.')) for row in self.rows(name))]
    def resolve_target(self,item):
        preferred=target(item)
        if preferred in self.names():return preferred
        return next((name for name in self.client_names() if normalized(name.removesuffix('行验货'))==normalized(item.get('customer'))),preferred)
    def profile(self,name):
        second=self.second(name);blocks=[];current=None;latest_date=None
        def styles(row):return {column_index_from_string(re.sub(r'\d','',c.get('r'))):c.get('s','0') for c in row.findall(tag('c'))}
        for row in sorted(self.rows(name),key=lambda r:int(r.get('r'))):
            vals=self.values(row)
            if isinstance(vals.get(1),(int,float)) and vals[1]>30000 and not any(normalized(vals.get(c)) for c in range(2,14)):
                latest_date=row;current=None
            elif normalized(vals.get(1))==normalized('Contract No.' if second else '客户'):
                current={'date':latest_date,'header':row,'data':[]};blocks.append(current)
            elif current is not None and vals.get(1 if second else 2) and vals.get(3 if second else 4):
                current['data'].append(row)
        # A fresh annual template may retain empty formatted rows instead of old orders.
        if not any(block['data'] for block in blocks):
            for block in blocks:
                header=int(block['header'].get('r'))
                blank=next((row for row in self.rows(name) if int(row.get('r'))>header and not any(self.values(row).values()) and row.findall(tag('c'))),None)
                if blank is not None:block['data']=[blank]
        valid=[block for block in blocks if block['data']]
        if not valid:raise HTTPException(422,'客户模板缺少有效表头或样式行：'+name)
        block=valid[-1];body={}
        for row in block['data']:body.update(styles(row))
        widths={i:13 for i in range(1,15)}
        cols=self.sheet(name).find(tag('cols'))
        if cols is not None:
            for col in cols:
                for i in range(int(col.get('min')),min(14,int(col.get('max')))+1):widths[i]=float(col.get('width','13'))
        return {'date':styles(block['date'] if block['date'] is not None else block['header']),
            'header':styles(block['header']),'data':body,'widths':widths,
            'templateRows':{'date':int(block['date'].get('r')) if block['date'] is not None else None,'header':int(block['header'].get('r')),'data':[int(r.get('r')) for r in block['data']]}}
    def row(self,name,number,values,styles,height=24):
        for existing in list(self.rows(name)):
            if int(existing.get('r'))==number:
                if any(normalized(v) for v in self.values(existing).values()):raise HTTPException(422,'写入位置与历史数据冲突：'+name)
                self.rows(name).remove(existing)
        row=E.Element(tag('row'),{'r':str(number),'ht':str(height),'customHeight':'1'})
        for col,value in enumerate(values,1):self.cell(row,col,value,styles.get(col,'0'))
        self.rows(name).append(row);return row
    def finish(self):
        defined=self.book.find(tag('definedNames'))
        if defined is not None:
            for name in defined:
                old=name.get('localSheetId')
                if old is not None and int(old)<len(self.original_order):
                    original=self.original_order[int(old)]
                    if original in self.names():name.set('localSheetId',str(self.names().index(original)))
        for name,node in getattr(self,'generated_names',[]):
            if defined is None:defined=E.SubElement(self.book,tag('definedNames'))
            node.set('localSheetId',str(self.names().index(name)));defined.append(node)
        for name,root in self.loaded.items():
            old=self.parts.get(self.path(name))
            if old and E.tostring(root)==E.tostring(E.fromstring(old)):continue
            rows=root.find(tag('sheetData'));rows[:]=sorted(rows,key=lambda r:int(r.get('r')))
            dim=root.find(tag('dimension'))
            if dim is not None:root.remove(dim) # optional; readers calculate from cells, never trust legacy whole-sheet dimensions
            # Keep OOXML worksheet child order valid after inserting merges and drawings.
            order=['sheetPr','dimension','sheetViews','sheetFormatPr','cols','sheetData','sheetCalcPr','sheetProtection','protectedRanges','scenarios','autoFilter','sortState','dataConsolidate','customSheetViews','mergeCells','phoneticPr','conditionalFormatting','dataValidations','hyperlinks','printOptions','pageMargins','pageSetup','headerFooter','rowBreaks','colBreaks','customProperties','cellWatches','ignoredErrors','smartTags','drawing','legacyDrawing','legacyDrawingHF','picture','oleObjects','controls','webPublishItems','tableParts','extLst']
            root[:]=sorted(root,key=lambda n:order.index(E.QName(n).localname) if E.QName(n).localname in order else len(order))
            self.parts[self.path(name)]=xml(root)
        self.parts['xl/workbook.xml']=xml(self.book);self.parts['xl/_rels/workbook.xml.rels']=xml(self.rels)
        self.parts['xl/styles.xml']=xml(self.styles);self.parts['[Content_Types].xml']=xml(self.types)
        result=BytesIO()
        with ZipFile(result,'w',ZIP_DEFLATED) as z:
            for name,content in self.parts.items():z.writestr(name,content)
        return result.getvalue()

def height(values,widths,p=None,styles=None):
    required=24
    for col,value in enumerate(values,1):
        size=p.font_size((styles or {}).get(col,'0')) if p else 11
        # Excel column units are digit widths (about 7px at the normal 11pt font),
        # not full CJK glyphs. Measure text in half-em units at the actual cell size.
        capacity=max(3,(widths.get(col,13)*7-3)/(size*2/3))
        lines=0
        for line in text(value).split('\n'):
            used=0;count=1
            for token in re.findall(r'[A-Za-z0-9_]+|[^A-Za-z0-9_]',line):
                length=sum(0 if unicodedata.combining(ch) else 2 if unicodedata.east_asian_width(ch) in 'WF' else 1.5 if ch in 'WM@%' else 0.55 if ch in 'il.,;:! ' else 1 for ch in token)
                if used and used+length>capacity:count+=1;used=0
                if length>capacity:
                    extra=int((length-0.001)//capacity);count+=extra;length-=extra*capacity
                used+=length
            lines+=count
        required=max(required,lines*(size*1.15)+6)
    return round(required,2)

def tall_row(p,name,number,values,styles,widths,meta):
    h=height(values,widths,p,styles);span=max(1,math.ceil(h/380))
    if number+span>1048576:raise HTTPException(413,'工作表已超出 Excel 行数限制')
    p.row(name,number,values,styles,h/span)
    if span>1:
        merges=p.sheet(name).find(tag('mergeCells'))
        if merges is None:merges=E.SubElement(p.sheet(name),tag('mergeCells'))
        for col in range(1,len(values)+1):
            ref=f'{get_column_letter(col)}{number}:{get_column_letter(col)}{number+span-1}'
            E.SubElement(merges,tag('mergeCell'),{'ref':ref});meta.setdefault('ownedMerges',{}).setdefault(name,[]).append(ref)
        for k in range(1,span):p.row(name,number+k,[],styles,h/span)
    return span

def historic_candidates(p,name,item):
    """Bind only exact date/customer/full identifiers; never fuzzy-match identity."""
    second=p.second(name);group='';found=[]
    for row in p.rows(name):
        vals=p.values(row);a=vals.get(1)
        if isinstance(a,(float,int)) and a>30000 and not any(normalized(vals.get(c)) for c in range(2,14)):
            try:group=from_excel(a,p.epoch).date().isoformat()
            except (ValueError,OverflowError):group=''
        po,sku=(vals.get(1),vals.get(3)) if second else (vals.get(2),vals.get(4))
        if normalized(po)!=normalized(item.get('poNo')) or normalized(sku)!=normalized(item.get('productNo')):continue
        if not second and normalized(vals.get(1))!=normalized(item.get('customer')):continue
        if group!=item['archiveDate']:continue
        found.append(row)
    if len(found)>1:raise HTTPException(422,f"记录 {item['id']} 在 {name} 有多个历史候选，未写入")
    if found:
        vals=p.values(found[0]);q=vals.get(5 if second else 6)
        if normalized(q)!=normalized(numbers(item)[0]):raise HTTPException(422,f"记录 {item['id']} 与历史候选数量不一致，未覆盖")
        return found[0]

def _build(content,previous,current,metadata,legacy=False):
    """previous/current keyed by stable entry ID; metadata is part of the version."""
    p=Package(content);meta=deepcopy(metadata);meta['formatVersion']=FORMAT_VERSION
    p.layouts={name:(2 if p.second(name) else 1) for name in p.client_names()};p.layouts.update(meta.get('clientLayouts',{}))
    meta.setdefault('ownedRows',{});meta.setdefault('bindings',{});meta.setdefault('weekSheets',{});meta.setdefault('detailSheets',[])
    for name,refs in meta.get('ownedLinks',{}).items():
        links=p.sheet(name).find(tag('hyperlinks'))
        if links is not None:
            for link in list(links):
                if link.get('ref') in refs:links.remove(link)
            if not len(links):p.sheet(name).remove(links)
    meta['ownedLinks']={};meta['summaryRows']={}
    # Remove only rows and merge ranges owned by this feature.
    for name,refs in meta.get('ownedMerges',{}).items():
        merges=p.sheet(name).find(tag('mergeCells'))
        if merges is not None:
            for merge in list(merges):
                if merge.get('ref') in refs:merges.remove(merge)
            if not len(merges):p.sheet(name).remove(merges)
    meta['ownedMerges']={}
    for name,coords in meta['ownedRows'].items():
        rows=p.rows(name)
        for row in list(rows):
            if int(row.get('r')) in coords:rows.remove(row)
    for identity,binding in meta['bindings'].items():
        if binding.get('original'):
            rows=p.rows(binding['sheet'])
            for row in list(rows):
                if int(row.get('r'))==binding['row']:rows.remove(row)
            item=current.get(identity,{})
            if item and p.resolve_target(item)==binding['sheet'] and item['archiveDate']==previous.get(identity,{}).get('archiveDate'):
                rows.append(E.fromstring(binding['original']))
    # Remove only resources reachable from the explicitly owned detail sheets.
    # A historical picture with a similar filename is not owned by this feature.
    remove_details(p,meta['detailSheets'])
    meta['detailSheets']=[];meta['ownedRows']={};newbindings={}
    profiles={name:p.profile(name) for name in p.client_names() if name not in meta.get('createdCustomers',[])}
    if not profiles:raise HTTPException(422,'底表没有可复用的客户样式')
    fallback=next((name for name in profiles if not p.second(name)),next(iter(profiles)))
    grouped=defaultdict(list);claimed=set()
    for identity,item in sorted(current.items()):
        name=p.resolve_target(item);old=meta['bindings'].get(identity)
        if name not in p.names() or name not in profiles:
            desired=p.second(name) if name in p.names() else name in SECOND
            source=next((candidate for candidate in profiles if p.second(candidate)==desired),None)
            if source is None:raise HTTPException(422,'最新底表缺少该客户适用的字段结构：'+name)
            if name not in p.names():
                p.add(name,source);p.layouts[name]=p.layouts[source]
                meta.setdefault('createdCustomers',[])
                if name not in meta['createdCustomers']:meta['createdCustomers'].append(name)
            elif name not in set(metadata.get('createdCustomers',[])):
                raise HTTPException(422,'未识别的客户表结构：'+name)
            profiles[name]=profiles[source]
        row=None
        if old and old.get('original') and old['sheet']==name and previous.get(identity,{}).get('archiveDate')==item['archiveDate']:
            row=E.fromstring(old['original'])
        elif identity not in previous:row=historic_candidates(p,name,item)
        if row is not None:
            coord=(name,int(row.get('r')))
            if coord in claimed:raise HTTPException(422,'多条系统记录对应同一人工历史行，未覆盖：'+name)
            claimed.add(coord);values=client_values(item,p.second(name))
            rows=p.rows(name)
            for r in list(rows):
                if int(r.get('r'))==coord[1]:rows.remove(r)
            original_styles={column_index_from_string(re.sub(r'\d','',c.get('r'))):c.get('s','0') for c in row.findall(tag('c'))}
            styles=data_styles(p,{'data':original_styles},p.second(name));h=height(values,profiles[name]['widths'],p,styles)
            if h>400:
                # Move this uniquely bound managed record to a full-height group at the end;
                # leave every other historical row/formula coordinate untouched.
                grouped[name].append((identity,item))
            else:
                edited=deepcopy(row);edited.set('ht',str(h));edited.set('customHeight','1')
                for col,value in enumerate(values,1):
                    oldcell=next((c for c in edited if c.get('r')==get_column_letter(col)+str(coord[1])),None)
                    independent={9:'inspectionFactory',10:'reportStatus',12:'responsible',13:'otherRemark'} if p.second(name) else {11:'firstResponsible',12:'secondResponsible',13:'otherRemark'}
                    if col in independent and independent[col] not in item:continue
                    if oldcell is not None and oldcell.find(tag('f')) is not None:continue
                    if oldcell is not None:edited.remove(oldcell)
                    p.cell(edited,col,value,styles.get(col,'0'))
                edited[:]=sorted(edited,key=lambda c:column_index_from_string(re.sub(r'\d','',c.get('r'))))
                rows.append(edited)
                newbindings[identity]={'sheet':name,'row':coord[1],'original':E.tostring(row,encoding='unicode')}
        else:grouped[name].append((identity,item))
    for name,items in grouped.items():
        profile=profiles[name];rows=p.rows(name)
        end=max([int(r.get('r')) for r in rows if any(normalized(v) for v in p.values(r).values())]+[0])+2
        owned=[];styles=data_styles(p,profile,p.second(name))
        dates=defaultdict(list)
        for identity,item in items:dates[displayed_date(item)].append((identity,item))
        for label,group in sorted(dates.items(),key=lambda pair:(pair[1][0][1]['archiveDate'],text(pair[0]))):
            ds={**profile['date'],1:p.date_style(profile['date'].get(1,'0'),profile['date'],label)}
            if not legacy:
                # Widen only this newly owned date header; historical columns and rows stay intact.
                widths={1:sum(profile['widths'].get(col,10) for col in (1,2,3))}
                p.row(name,end,[label],ds,height([label],widths,p,ds))
                merges=p.sheet(name).find(tag('mergeCells'))
                if merges is None:merges=E.SubElement(p.sheet(name),tag('mergeCells'))
                ref=f'A{end}:C{end}'
                E.SubElement(merges,tag('mergeCell'),{'ref':ref})
                meta.setdefault('ownedMerges',{}).setdefault(name,[]).append(ref)
            else:p.row(name,end,[label],ds,height([label],profile['widths'],p,ds))
            owned.append(end);end+=1
            p.row(name,end,HEAD2 if p.second(name) else HEAD1,profile['header']);owned.append(end);end+=1
            for identity,item in sorted(group,key=record_order):
                values=client_values(item,p.second(name));span=tall_row(p,name,end,values,styles,profile['widths'],meta)
                owned.extend(range(end,end+span))
                newbindings[identity]={'sheet':name,'row':end};end+=span
            end+=1
        meta['ownedRows'][name]=owned
    meta['bindings']=newbindings
    if not legacy:
        meta['clientLayouts']=p.layouts
        return finish_details(p,current,meta)
    weeks=defaultdict(list)
    for identity,item in current.items():weeks[item['week']].append((identity,item))
    for week,items in sorted(weeks.items()):
        name=meta['weekSheets'].get(week)
        if not name:
            last=(date.fromisoformat(week)+timedelta(days=6)).isoformat()
            preferred=week+'至'+last[5:];name=preferred;suffix=1
            while name in p.names():
                name=preferred+'（续'+str(suffix)+'）';suffix+=1
            p.add(name,'ZURU行验货');meta['weekSheets'][week]=name
        rows=p.rows(name);rows[:]=[]
        profile=profiles['ZURU行验货'];widths=dict(enumerate([26,18,26,24,26,30,12,12,12,18,16,26,22,65],1))
        root=p.sheet(name);cols=root.find(tag('cols'))
        if cols is not None:root.remove(cols)
        cols=E.Element(tag('cols'))
        for c,w in widths.items():E.SubElement(cols,tag('col'),{'min':str(c),'max':str(c),'width':str(w),'customWidth':'1'})
        root.insert(list(root).index(rows),cols)
        mapping={1:1,2:1,3:2,4:3,5:4,6:5,7:6,8:7,9:8,10:9,11:9,12:2,13:9,14:10}
        header={c:profile['header'].get(old,'0') for c,old in mapping.items()};styles={c:p.style(profile['data'].get(old,'0')) for c,old in mapping.items()}
        p.row(name,1,SUMMARY,header,30)
        i=2;dates=defaultdict(list)
        for identity,item in sorted(items,key=record_order):dates[item['archiveDate']].append((identity,item))
        for day,day_items in dates.items():
            label=date.fromisoformat(day);ds={**profile['date'],1:p.date_style(profile['date'].get(1,'0'),profile['date'],label)}
            p.row(name,i,[label],ds);i+=1
            customers=defaultdict(list)
            for identity,item in day_items:customers[normalized(item['customer'])].append((identity,item))
            for group in customers.values():
                for identity,item in group:
                    cp=profiles[p.resolve_target(item)];second=p.resolve_target(item) in SECOND
                    mp={1:1,2:2,3:1,4:2,5:3,6:4,7:5,8:6,9:7,10:9,11:10,12:1,13:8,14:11} if second else mapping
                    styles={col:p.style(cp['data'].get(old,'0'),col in {1,2,3,4,5,6,10,11,12,14}) for col,old in mp.items()}
                    displayed=displayed_date(item);styles[1]=p.date_style(styles[1],cp['date'],displayed)
                    q,u,c=numbers(item)
                    values=[displayed,item.get('customer'),item.get('poNo'),item.get('country'),item.get('productNo'),item.get('productName'),q,u,c,item.get('production'),item.get('agency'),item.get('customerPo'),item.get('result',''),item.get('remark')]
                    meta['summaryRows'][identity]={'sheet':name,'row':i}
                    i+=tall_row(p,name,i,values,styles,widths,meta)
                amounts=[numbers(item) for _,item in group]
                p.row(name,i,['客户小计',group[0][1]['customer'],'','','','',sum(v[0] for v in amounts if isinstance(v[0],(int,float))),'',sum(v[2] for v in amounts if isinstance(v[2],(int,float)))],header);i+=1
            # Sum source detail rows directly, never already calculated customer subtotal rows.
            amounts=[numbers(item) for _,item in day_items]
            p.row(name,i,['当日小计','','','','','',sum(v[0] for v in amounts if isinstance(v[0],(int,float))),'',sum(v[2] for v in amounts if isinstance(v[2],(int,float)))],header);i+=1
    # A moved record leaves an existing empty weekly summary, preserving other history and sheet identity.
    for week,name in meta['weekSheets'].items():
        if week not in weeks:p.rows(name)[:]=[];p.row(name,1,SUMMARY,profiles['ZURU行验货']['header'],30)
    add_details(p,current,meta,profiles['ZURU行验货'])
    for identity,detail_row in meta.get('detailRows',{}).items():
        for binding,col in [(meta['bindings'][identity],1 if p.second(meta['bindings'][identity]['sheet']) else 2),(meta['summaryRows'][identity],3)]:
            name=binding['sheet'];ref=get_column_letter(col)+str(binding['row']);root=p.sheet(name);links=root.find(tag('hyperlinks'))
            if links is None:links=E.SubElement(root,tag('hyperlinks'))
            if any(link.get('ref')==ref for link in links):continue # keep an existing manual hyperlink
            E.SubElement(links,tag('hyperlink'),{'ref':ref,'location':"'累计验货详情'!A"+str(detail_row),'tooltip':'查看该笔验货完整文字'})
            meta['ownedLinks'].setdefault(name,[]).append(ref)
    ordered=[meta['weekSheets'][w] for w in sorted(meta['weekSheets'],reverse=True)]
    sheet_nodes=list(p.sheets);p.sheets[:]=sorted(sheet_nodes,key=lambda s:(0,ordered.index(s.get('name'))) if s.get('name') in ordered else (1,sheet_nodes.index(s)))
    if ordered:
        for view in p.book.findall(tag('bookViews')+'/'+tag('workbookView')):view.set('activeTab','0');view.set('firstSheet','0')
    # Reset only generated views; keep all original sheet navigation settings.
    for name in meta.get('createdCustomers',[])+list(meta['weekSheets'].values())+meta['detailSheets']:
        root=p.sheet(name);old=root.find(tag('sheetViews'))
        if old is not None:root.remove(old)
        view=E.SubElement(E.SubElement(root,tag('sheetViews')),tag('sheetView'),{'workbookViewId':'0','topLeftCell':'A1'})
        E.SubElement(view,tag('selection'),{'activeCell':'A1','sqref':'A1'})
    return p.finish(),meta

def remove_details(p,names):
    def rel_path(owner):return posixpath.dirname(owner)+'/_rels/'+posixpath.basename(owner)+'.rels'
    def references(owner):
        raw=p.parts.get(rel_path(owner))
        if not raw:return []
        return [(r.get('Type'),posixpath.normpath(posixpath.join(posixpath.dirname(owner),r.get('Target'))).lstrip('/') if not r.get('Target').startswith('/') else r.get('Target').lstrip('/'))
            for r in E.fromstring(raw) if r.get('TargetMode')!='External']
    roots={p.path(name) for name in names if name in p.names()};resources=set();pending=list(roots)
    while pending:
        for kind,target in references(pending.pop()):
            if kind in {R+'/drawing',R+'/image'} and target not in resources:resources.add(target);pending.append(target)
    owned_rels={rel_path(owner) for owner in roots|resources};protected=set()
    for path,raw in p.parts.items():
        if not path.endswith('.rels') or path in owned_rels:continue
        owner=posixpath.join(posixpath.dirname(posixpath.dirname(path)),posixpath.basename(path)[:-5])
        protected.update(target for _,target in references(owner) if target in resources)
    pending=list(protected)
    while pending:
        for _,target in references(pending.pop()):
            if target in resources and target not in protected:protected.add(target);pending.append(target)
    removed=resources-protected
    for resource in removed:p.parts.pop(resource,None);p.parts.pop(rel_path(resource),None)
    for entry in list(p.types):
        if entry.get('PartName','').lstrip('/') in removed:p.types.remove(entry)
    for name in names:
        if name in p.names():p.remove(name)


def add_details(p,current,meta,profile):
    """Text-only traceability. No image table, bytes, URL, placeholder or photo-height row."""
    if not current:return
    meta['detailRows']={};meta['detailSpans']={}
    name='累计验货详情'
    if name in p.names():raise HTTPException(422,'详情表名与人工历史冲突')
    p.add(name,'ZURU行验货');meta['detailSheets'].append(name)
    headers=['业务日期／排期','客户','PO#','货号','验货结果','现场原文','其他备注','记录编号','源版本','数量','装箱','箱数']
    independent=any('inspectionProblem' in item for item in current.values())
    if independent:headers.append('独立验货问题')
    p.row(name,1,headers,profile['header'],30)
    root=p.sheet(name);cols=root.find(tag('cols'))
    if cols is not None:root.remove(cols)
    cols=E.Element(tag('cols'));widths=dict(enumerate([26,18,28,28,12,60,40,38,14,12,12,12]+([60] if independent else []),1))
    for c,w in widths.items():E.SubElement(cols,tag('col'),{'min':str(c),'max':str(c),'width':str(w),'customWidth':'1'})
    root.insert(list(root).index(root.find(tag('sheetData'))),cols)
    mapping={1:1,2:1,3:2,4:4,5:9,6:10,7:13,8:2,9:8,10:6,11:7,12:8,13:10}
    styles={c:p.style(profile['data'].get(old,'0'),c in {1,2,3,4,6,7,8,13}) for c,old in mapping.items()}
    rownum=2
    for identity,item in sorted(current.items(),key=record_order):
        meta['detailRows'][identity]=rownum
        fields=[text(item.get(key)) for key in ('remark','otherRemark','inspectionProblem')]
        chunks=[[value[i:i+500] for i in range(0,len(value),500)] or [''] for value in fields]
        for index in range(max(map(len,chunks))):
            displayed=displayed_date(item);styles[1]=p.date_style(styles[1],profile['date'],displayed)
            texts=[parts[index] if index<len(parts) else '' for parts in chunks]
            values=[displayed,item.get('customer'),item.get('poNo'),item.get('productNo'),item.get('result',''),texts[0],texts[1],identity,item.get('revision'),*numbers(item)]
            if independent:values.append(texts[2])
            rownum+=tall_row(p,name,rownum,values,styles,widths,meta)
        meta['detailSpans'][identity]=rownum-meta['detailRows'][identity]


def sheet_values(p,name):
    return [(c.get('r'),p.values(row).get(column_index_from_string(re.sub(r'\d','',c.get('r')))),E.tostring(c.find(tag('f')),encoding='unicode') if c.find(tag('f')) is not None else None)
            for row in p.rows(name) for c in row if c.find(tag('f')) is not None or p.values(row).get(column_index_from_string(re.sub(r'\d','',c.get('r')))) not in ('',None)]

def sheet_digest(p,name):
    from hashlib import sha256
    return sha256(json.dumps(sheet_values(p,name),ensure_ascii=False,sort_keys=True,default=str).encode()).hexdigest()

def build(content,previous,current,metadata):
    # Validate owned pages before any regeneration, so a manual unique note cannot disappear.
    original=Package(content);owned=set(metadata.get('weekSheets',{}).values())|set(metadata.get('detailSheets',[]))|({metadata['detailSheet']} if metadata.get('detailSheet') else set())
    for old in owned:
        if old not in original.names():continue
        path=original.path(old);folder,filename=path.rsplit('/',1);relpath=folder+'/_rels/'+filename+'.rels'
        if relpath in original.parts and len(E.fromstring(original.parts[relpath])):
            raise HTTPException(409,'自动页有额外批注、图片或链接，已保留：'+old)
    for defined in original.book.findall('.//'+tag('definedName')):
        local=defined.get('localSheetId');local_name=original.names()[int(local)] if local is not None and int(local)<len(original.names()) else None
        if local_name in owned and defined.get('name') in ('_xlnm.Print_Area','_xlnm.Print_Titles','_xlnm._FilterDatabase'):continue
        if any(old+'!' in (defined.text or '') or "'"+old+"'!" in (defined.text or '') for old in owned):
            raise HTTPException(409,'名称定义引用自动页，已保留原文件')
    for name in original.names():
        if name in owned:continue
        for link in original.sheet(name).findall('.//'+tag('hyperlink')):
            if link.get('ref') in metadata.get('ownedLinks',{}).get(name,[]):continue
            if any(old+'!' in link.get('location','') or "'"+old+"'!" in link.get('location','') for old in owned):
                raise HTTPException(409,'历史链接引用自动页，已保留：'+name)
        formulas=[f.text or '' for f in original.sheet(name).findall('.//'+tag('f'))]
        if any(old+'!' in formula or "'"+old+"'!" in formula for formula in formulas for old in owned):raise HTTPException(409,'历史公式引用自动页，未替换：'+name)

    if metadata.get('detailSheet'):
        p=Package(content);name=metadata['detailSheet']
        if name not in p.names() or sheet_digest(p,name)!=metadata.get('detailDigest'):raise HTTPException(409,'自动生成的验货明细页含未登记修改，已保留原文件，请先核对')
    elif metadata.get('weekSheets') or metadata.get('detailSheets'):
        p=Package(content);expected,_=_build(content,previous,previous,metadata,legacy=True);q=Package(expected)
        for name in [*metadata.get('weekSheets',{}).values(),*metadata.get('detailSheets',[])]:
            if name not in p.names() or name not in q.names() or sheet_values(p,name)!=sheet_values(q,name):raise HTTPException(409,'旧自动页含独有内容，未删除：'+name)
    return _build(content,previous,current,metadata)

def finish_details(p,current,meta):
    from app.services.qc_inspection_sheet import export_records
    from app.services.qc_workbook_style import copy_detail_sheet
    obsolete=set(meta.get('weekSheets',{}).values())|set(meta.get('detailSheets',[]))|({meta['detailSheet']} if meta.get('detailSheet') else set())
    for name in p.names():
        if name in obsolete:continue
        for formula in p.sheet(name).findall('.//'+tag('f')):
            if any(old+'!' in (formula.text or '') or "'"+old+"'!" in (formula.text or '') for old in obsolete):raise HTTPException(409,'历史公式引用旧自动页，未删除：'+name)
    for name in obsolete:
        if name in p.names():p.remove(name)
    meta['weekSheets']={};meta['detailSheets']=[];meta['detailRows']={};meta['detailSpans']={};meta['summaryRows']={}
    ids=meta.get('detailIds',list(current));items=[current[i] for i in ids if i in current]
    name=meta.get('detailSheet','验货明细');preferred=name;n=1
    while name in p.names():name=preferred+'（'+str(n)+'）';n+=1
    copy_detail_sheet(p,export_records(items,include_photos=False,epoch=p.epoch),name)
    meta['detailSheet']=name;meta['detailIds']=[item['id'] for item in items]
    meta['summaryRows']={item['id']:{'sheet':name,'row':i} for i,item in enumerate(items,2)}
    node=next(s for s in p.sheets if s.get('name')==name);p.sheets.remove(node);p.sheets.insert(0,node)
    for view in p.book.findall(tag('bookViews')+'/'+tag('workbookView')):view.set('activeTab','0');view.set('firstSheet','0')
    # Only loaded pages that actually changed are serialized by Package; original values remain untouched.
    meta['detailDigest']=sheet_digest(p,name)
    return p.finish(),meta
