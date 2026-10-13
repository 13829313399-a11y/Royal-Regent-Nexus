"""Carton-style location setup using fabric evidence, identity and factory locks."""
import json
from io import BytesIO
from uuid import uuid4, uuid5
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import inspect, select
from app.core.time import business_now
from app.models.fabric_master import FabricMasterRecord, FabricMasterChange
from app.schemas.fabric_master import LocationPreviewRequest, LocationInput
from app.services import fabric_master as master
from app.services import fabric_procurement as source
from app.services.fabric_procurement_parser import digest, dump
from app.services.carton_master_import import location_bins, parse, template as carton_template


def template():
    # Reuse only the empty workbook shape; fabric explanatory contract is separate.
    book = load_workbook(BytesIO(carton_template("locations")))
    notes = book["说明与示例"]
    book.remove(notes)
    notes = book.create_sheet("说明与示例")
    for message in (
        "仅导入“导入数据”页，说明页不导入。支持 xlsx/xlsm，5 MB、1000 行，展开后最多 1000 个仓位。",
        "两列均填真实文本，不接受公式或数字单元格。仓位可填 A01-A03、B01 或列表，保留前导零。",
        "仓位编码在布料仓中唯一。不同仓库不能重复使用同一编码；可使用 A区01、B区01。",
        "预览核对后保存；已有同仓仓位跳过，停用仓位不会自动启用，跨仓同编码报错。",
        "新增仓位默认启用；首个仓位建立仓库。导入不增加库存，也不更改历史单据。",
    ):
        notes.append([message])
    notes.append(["仓库", "仓位或范围"])
    notes.append(["布料一仓", "A01-A03"])
    notes.column_dimensions["A"].width = 110
    notes.column_dimensions["B"].width = 24
    output = BytesIO(); book.save(output); book.close()
    return output.getvalue()


def file_rows(content, filename):
    entries, errors = parse(content, filename, "locations")
    if errors:
        raise HTTPException(422, "；".join(errors[:20]))
    return [LocationInput(warehouse=entry["warehouse"], bins=code) for _, entry in entries for code in entry["bins"]]


def locations(db, include_deleted=False):
    return [row for row in master.records(db) if row.kind == "LOCATION" and (include_deleted or not master.is_deleted(row))]


def request_data(payload):
    # Preserve signatures for legacy requests without a category.
    body = payload.model_dump(mode="json")
    for row in body.get("rows", []):
        if not row.get("category"):
            row.pop("category", None)
    return body


def preview(db, actor, payload):
    master.ensure_schema(db)
    current = locations(db, include_deleted=True)
    by_code = {}
    warehouse_names = {json.loads(row.data_json).get("warehouse", "") for row in current}
    for row in current:
        by_code.setdefault(row.code.casefold(), []).append(row)
    items, errors, seen = [], [], set()
    for index, row in enumerate(payload.rows, 1):
        try:
            if any(c in row.warehouse for c in ",，、;；\r\n|"):
                raise ValueError("仓库名称每项只填写一个")
            matches = [name for name in warehouse_names if name.casefold() == row.warehouse.casefold()]
            if len(matches) > 1:
                raise ValueError("仓库名称有多个大小写对应，请先核对旧资料")
            warehouse = matches[0] if matches else row.warehouse
            warehouse_names.add(warehouse)
            bins = location_bins(row.bins, 1000 - len(items))
            for code in bins:
                key = code.casefold()
                if key in seen:
                    raise ValueError(f"仓位 {code} 在本次资料中重复")
                seen.add(key)
                existing = by_code.get(key, [])
                if len(existing) > 1:
                    raise ValueError(f"仓位 {code} 有多个旧编码对应，请先核对")
                prior = existing[0] if existing else None
                if prior and master.is_deleted(prior):
                    raise ValueError(f"仓位 {code} 已删除，请先从已删除仓位中恢复，不能重复建立编码")
                if prior and json.loads(prior.data_json).get("warehouse") != warehouse:
                    raise ValueError(f"仓位 {code} 已属于其他仓库，不能通过导入移动")
                category = json.loads(prior.data_json).get("category", "") if prior else row.category
                if prior and row.category and category != row.category:
                    raise ValueError(f"仓位 {code} 已有不同的原料分区，请先在仓位资料中核对修改")
                items.append({"warehouse": warehouse, "code": code, "action": "UNCHANGED" if prior else "NEW",
                              "status": prior.status if prior else "ACTIVE", "category": category})
        except ValueError as exc:
            errors.append(f"第 {index} 项：{exc}")
    token = digest(["fabric-locations-v1", actor.id, request_data(payload),
                    sorted([master.record_data(row) for row in current], key=lambda row: row["id"])])
    return {"items": items, "errors": errors, "new": sum(row["action"] == "NEW" for row in items),
            "unchanged": sum(row["action"] == "UNCHANGED" for row in items), "preview_token": token, "stock_posted": False}


def replay(db, actor, payload):
    request_hash = digest([actor.id, request_data(payload)])
    previous = db.scalar(select(FabricMasterChange).where(FabricMasterChange.factory_id == source.FACTORY,
                                                         FabricMasterChange.request_id == str(payload.request_id)))
    if previous:
        operation = json.loads(previous.after_json).get("location_operation")
        if previous.request_hash != request_hash or not operation:
            raise HTTPException(409, "该请求已用于其他资料或账号，请刷新核对")
        db.rollback()
        return request_hash, operation
    return request_hash, None


def write_operation(db, actor, payload, request_hash, changes, result, state):
    request_ids = [str(payload.request_id)] + [str(uuid5(payload.request_id, f"fabric-location-{index}")) for index in range(1, len(changes))]
    for start in range(0, len(request_ids), 500):
        if db.scalar(select(FabricMasterChange.id).where(FabricMasterChange.factory_id == source.FACTORY,
                                                        FabricMasterChange.request_id.in_(request_ids[start:start + 500]))):
            raise HTTPException(409, "请求标识已用于其他资料，请刷新核对")
    now = business_now().isoformat()
    for index, (record, before) in enumerate(changes):
        record.revision += 1; record.updated_at = now
        db.add(record); db.flush()
        after = master.record_data(record)
        if index == 0:
            after["location_operation"] = result
        db.add(FabricMasterChange(id=uuid4().hex, factory_id=source.FACTORY, record_id=record.id,
            request_id=request_ids[index], request_hash=request_hash, before_json=dump(before), after_json=dump(after),
            actor_id=actor.id, actor_name=actor.display_name, occurred_at=now))
    state.revision += 1
    db.commit()
    return result


def apply(db, actor, payload):
    master.ensure_schema(db)
    state = source.lock_factory(db)
    request_hash, previous = replay(db, actor, payload)
    if previous is not None:
        return previous
    plan = preview(db, actor, LocationPreviewRequest(factory_id=payload.factory_id, rows=payload.rows))
    if plan["preview_token"] != payload.preview_token:
        raise HTTPException(409, "仓位资料已变化，请重新预览")
    if plan["errors"]:
        raise HTTPException(422, "；".join(plan["errors"][:20]))
    if not plan["new"]:
        raise HTTPException(422, "所有仓位均已存在，无需重复保存")
    changes = []
    for item in plan["items"]:
        if item["action"] == "NEW":
            record = FabricMasterRecord(id=uuid4().hex, factory_id=source.FACTORY, kind="LOCATION", code=item["code"],
                name=item["code"], status="ACTIVE", revision=0,
                data_json=dump({"warehouse": item["warehouse"], "category": item["category"], "note": ""}))
            changes.append((record, {}))
    result = {"new": plan["new"], "unchanged": plan["unchanged"], "stock_posted": False}
    return write_operation(db, actor, payload, request_hash, changes, result, state)


def rename(db, actor, payload):
    master.ensure_schema(db)
    state = source.lock_factory(db)
    request_hash, previous = replay(db, actor, payload)
    if previous is not None:
        return previous
    if payload.warehouse == payload.name or any(c in payload.name for c in ",，、;；\r\n|"):
        raise HTTPException(422, "请填写不同的单一仓库名称")
    current = locations(db)
    group = [row for row in current if json.loads(row.data_json).get("warehouse") == payload.warehouse]
    if not group or master.warehouse_tokens(current).get(payload.warehouse) != payload.expected_group_token:
        raise HTTPException(409, "仓库或仓位已变化，请刷新后修改")
    if any(json.loads(row.data_json).get("warehouse", "").casefold() == payload.name.casefold() for row in current if row not in group):
        raise HTTPException(409, "目标仓库名称已存在，请使用不同名称")
    changes = []
    for row in sorted(group, key=lambda row: row.id):
        before = master.record_data(row)
        data = json.loads(row.data_json); data["warehouse"] = payload.name
        row.data_json = dump(data); changes.append((row, before))
    return write_operation(db, actor, payload, request_hash, changes,
                           {"warehouse": payload.name, "updated": len(group), "stock_posted": False}, state)


def used_locations(db, selected):
    """Include immutable and reversed history, not just today's net stock."""
    from app.models.fabric_receiving import FabricReceipt, FabricStockBatch
    from app.models.warehouse_operations import WarehouseOperation
    identities = {row.id for row in selected}
    aliases = {}

    def remember(value):
        identifier = value.get('id')
        if identifier not in identities:
            return
        code = value.get('code', '')
        warehouse = value.get('data', {}).get('warehouse', '')
        for name in (code, *(warehouse + separator + code for separator in ('／', '/', ' / ', ' ／ ', ' '))):
            if name:
                aliases.setdefault(name.casefold(), set()).add(identifier)

    for row in selected:
        remember(master.record_data(row))
    for change in db.scalars(select(FabricMasterChange).where(FabricMasterChange.factory_id == source.FACTORY)):
        if change.record_id in identities:
            remember(json.loads(change.before_json)); remember(json.loads(change.after_json))
    used = set()

    def scan(value, key=''):
        if isinstance(value, dict):
            for name, part in value.items():
                scan(part, name)
        elif isinstance(value, list):
            for part in value:
                scan(part, key)
        elif isinstance(value, str):
            if value in identities:
                used.add(value)
            if key in {'location', 'location_label'}:
                used.update(aliases.get(value.strip().casefold(), set()))

    tables = set(inspect(db.get_bind()).get_table_names())
    if FabricStockBatch.__tablename__ in tables:
        for location in db.scalars(select(FabricStockBatch.location).where(FabricStockBatch.factory_id == source.FACTORY)):
            scan(location, 'location')
    if FabricReceipt.__tablename__ in tables:
        for payload in db.scalars(select(FabricReceipt.payload_json).where(FabricReceipt.factory_id == source.FACTORY)):
            scan(json.loads(payload))
    if WarehouseOperation.__tablename__ in tables:
        for payload in db.scalars(select(WarehouseOperation.payload_json).where(WarehouseOperation.factory_id == source.FACTORY, WarehouseOperation.warehouse == 'fabric')):
            scan(json.loads(payload))
    return used


def manage(db, actor, payload):
    master.ensure_schema(db)
    state = source.lock_factory(db)
    request_hash, previous = replay(db, actor, payload)
    if previous is not None:
        return previous
    if payload.action != 'CATEGORY' and payload.category:
        raise HTTPException(422, '删除或恢复仓位时不能同时变更原料分区')
    by_id = {row.id: row for row in locations(db, include_deleted=True)}
    selected = []
    for item in payload.selections:
        row = by_id.get(item.id)
        if not row or row.revision != item.expected_revision or master.is_deleted(row) != (payload.action == 'RESTORE'):
            raise HTTPException(409, '所选仓位已变化，请刷新后重新选择')
        selected.append(row)
    if len({row.id for row in selected}) != len(selected):
        raise HTTPException(422, '不能重复选择同一仓位')
    if payload.action == 'DELETE':
        used = used_locations(db, selected)
        if used:
            names = '、'.join(row.code for row in selected if row.id in used)
            raise HTTPException(409, f'仓位 {names} 已有库存或业务记录，不能删除；请取消勾选这些仓位，需停止使用时可改为停用。')
    changes = []
    for row in selected:
        before = master.record_data(row)
        data = json.loads(row.data_json)
        if payload.action == 'CATEGORY':
            data['category'] = payload.category
        elif payload.action == 'DELETE':
            data['_deleted'] = {'previous_status': row.status, 'at': business_now().isoformat()}
            row.status = 'INACTIVE'
        else:
            row.status = data.pop('_deleted')['previous_status']
        row.data_json = dump(data)
        changes.append((row, before))
    return write_operation(db, actor, payload, request_hash, changes,
                           {'action': payload.action, 'updated': len(selected), 'stock_posted': False}, state)
