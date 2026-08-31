from __future__ import annotations

from datetime import date
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import and_, select
from sqlalchemy.orm import Session

from app.models.injection_schedule import (
    InjectionScheduleLine,
    InjectionScheduleMachine,
    InjectionScheduleOrderDemand,
    InjectionScheduleShiftOutput,
)
from app.services.injection_scheduling.read_service import list_orders

MAX_EXPORT_ROWS = 10_000
TABLE_COLUMNS = {
    "order_no": "单号",
    "product_code": "货号",
    "product_name": "名称",
    "mold_code": "工模",
    "priority": "优先级",
    "status": "状态",
    "quantity_sets": "数量(套)",
    "total_sets": "总套数",
    "order_shots": "订单啤数",
    "qualified_shots": "已啤数",
    "remaining_shots": "欠数",
    "completion_percent": "完成率(%)",
    "delivery_due_date": "交货期",
    "warehouse": "仓库",
    "color": "颜色",
    "pigment_code": "色粉号",
    "material_name": "用料",
    "material_status": "材料状态",
    "required_machine_a_label": "必需机安",
    "daily_target": "日产量",
    "material_weight_kg": "用料重(kg)",
    "data_completeness_status": "数据完整性",
    "remark": "备注",
}
DEFAULT_COLUMNS = list(TABLE_COLUMNS)


def _styled_workbook(title: str) -> tuple[Workbook, object]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.freeze_panes = "A2"
    return workbook, sheet


def _finish(workbook: Workbook, sheet, row_count: int, column_count: int) -> bytes:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(horizontal="center")
    if row_count:
        sheet.auto_filter.ref = (
            f"A1:{sheet.cell(row=1, column=column_count).column_letter}{row_count + 1}"
        )
    for column in range(1, column_count + 1):
        values = [
            str(sheet.cell(row=row, column=column).value or "")
            for row in range(1, min(row_count + 2, 300))
        ]
        sheet.column_dimensions[
            sheet.cell(row=1, column=column).column_letter
        ].width = min(max(max(map(len, values), default=8) + 2, 10), 36)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def export_table(
    db: Session,
    *,
    factory_id: str,
    columns: list[str],
    search: str,
    status: str,
    priority: str,
) -> bytes:
    selected = columns or DEFAULT_COLUMNS
    invalid = [column for column in selected if column not in TABLE_COLUMNS]
    if invalid or len(selected) > 60:
        raise HTTPException(status_code=422, detail=f"导出字段不受支持：{invalid}")
    rows = list_orders(db, factory_id, search=search, status=status, priority=priority)
    if len(rows) > MAX_EXPORT_ROWS:
        raise HTTPException(
            status_code=413, detail=f"当前筛选超过 {MAX_EXPORT_ROWS} 行，请缩小范围"
        )
    workbook, sheet = _styled_workbook("当前排产台账")
    sheet.append([TABLE_COLUMNS[column] for column in selected])
    for row in rows:
        sheet.append([row.get(column) for column in selected])
    return _finish(workbook, sheet, len(rows), len(selected))


def export_shift_matrix(
    db: Session, *, factory_id: str, start_date: str, end_date: str
) -> bytes:
    try:
        start, end = date.fromisoformat(start_date), date.fromisoformat(end_date)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="导出日期格式不正确") from exc
    if end < start or (end - start).days > 30:
        raise HTTPException(status_code=422, detail="白夜班矩阵导出范围最多 31 天")
    rows = db.execute(
        select(
            InjectionScheduleMachine.machine_code,
            InjectionScheduleOrderDemand.order_no,
            InjectionScheduleOrderDemand.product_code,
            InjectionScheduleOrderDemand.mold_code,
            InjectionScheduleShiftOutput.production_date,
            InjectionScheduleShiftOutput.shift,
            InjectionScheduleShiftOutput.qualified_shots,
            InjectionScheduleShiftOutput.defect_shots,
            InjectionScheduleShiftOutput.downtime_minutes,
        )
        .join(
            InjectionScheduleLine,
            and_(
                InjectionScheduleLine.id
                == InjectionScheduleShiftOutput.schedule_line_id,
                InjectionScheduleLine.factory_id
                == InjectionScheduleShiftOutput.factory_id,
            ),
        )
        .join(
            InjectionScheduleOrderDemand,
            and_(
                InjectionScheduleOrderDemand.id
                == InjectionScheduleLine.order_demand_id,
                InjectionScheduleOrderDemand.factory_id
                == InjectionScheduleLine.factory_id,
            ),
        )
        .outerjoin(
            InjectionScheduleMachine,
            and_(
                InjectionScheduleMachine.id == InjectionScheduleLine.machine_id,
                InjectionScheduleMachine.factory_id == InjectionScheduleLine.factory_id,
            ),
        )
        .where(
            InjectionScheduleShiftOutput.factory_id == factory_id,
            InjectionScheduleShiftOutput.production_date >= start_date,
            InjectionScheduleShiftOutput.production_date <= end_date,
        )
        .order_by(
            InjectionScheduleShiftOutput.production_date,
            InjectionScheduleMachine.machine_code,
            InjectionScheduleShiftOutput.shift,
        )
        .limit(MAX_EXPORT_ROWS + 1)
    ).all()
    if len(rows) > MAX_EXPORT_ROWS:
        raise HTTPException(status_code=413, detail="白夜班矩阵超过导出行数上限")
    workbook, sheet = _styled_workbook("白夜班产量矩阵")
    headers = [
        "机号",
        "单号",
        "货号",
        "工模",
        "生产日期",
        "班次",
        "合格啤数",
        "不良啤数",
        "停机分钟",
    ]
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row))
    return _finish(workbook, sheet, len(rows), len(headers))
