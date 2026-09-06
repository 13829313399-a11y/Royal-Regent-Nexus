"""Authoritative server values, scoped SQL results and a stable exchange contract."""

import io
import json
from datetime import date, timedelta
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select

from app.models.injection_scheduling import (
    FactorySettings,
    HistoricalOutput,
    Machine,
    ReportAllocation,
    Run,
    RunDemand,
    ShiftReport,
)

from .calculations import decimal, now, quantities, timestamp
from .common import jsonable, record
from .field_registry import FACTORIES, FIELDS, LEGACY_FIELDS
from .queries import ordered_query, revision

SCHEMA = "injection-v3-exchange-1"


def export_plan(db, payload):
    start = date.fromisoformat(payload.window_start)
    end = start + timedelta(days=payload.window_days)
    rows = list(db.scalars(ordered_query(payload)))
    ids = [d.id for d in rows]
    shifts = {d: {} for d in ids}
    report_metadata = {d: {} for d in ids}
    before = {d.id: d.opening_shots for d in rows}
    before_units = dict.fromkeys(ids, decimal(0))
    historical_keys = {d: set() for d in ids}
    for h in db.scalars(
        select(HistoricalOutput).where(
            HistoricalOutput.demand_id.in_(ids), HistoricalOutput.production_date < end
        )
    ):
        if not h.counts_toward_demand:
            continue
        if h.production_date < start:
            before[h.demand_id] += h.quantity
            continue
        key = f"{h.production_date.isoformat()}|{h.shift_code}"
        historical_keys[h.demand_id].add(key)
        shifts[h.demand_id][key] = shifts[h.demand_id].get(key, 0) + h.quantity
    allocated = db.execute(
        select(RunDemand, ReportAllocation, ShiftReport)
        .join(ReportAllocation, ReportAllocation.run_demand_id == RunDemand.id)
        .join(ShiftReport, ShiftReport.id == ReportAllocation.report_id)
        .where(RunDemand.demand_id.in_(ids), ShiftReport.production_date < end)
    ).all()
    for link, allocation, report in allocated:
        if report.production_date < start:
            if allocation.unit == "SHOT":
                before[link.demand_id] += allocation.quantity
            else:
                before_units[link.demand_id] += allocation.quantity
            continue
        key = f"{report.production_date.isoformat()}|{report.shift_code}"
        # Each row declares its unit. Co-output good pieces are never added to M.
        shifts[link.demand_id][key] = (
            shifts[link.demand_id].get(key, 0) + allocation.quantity
        )
        report_metadata[link.demand_id].setdefault(key, []).append(
            {
                "id": report.id,
                "run_id": report.run_id,
                "revision": report.revision,
                "report_key": {
                    "run_id": report.run_id,
                    "production_date": report.production_date.isoformat(),
                    "shift_code": report.shift_code,
                    "segment_key": report.segment_key,
                },
                "segment_key": report.segment_key,
                "physical_shots": float(report.physical_shots),
                "allocated": float(allocation.quantity),
                "unit": allocation.unit,
            }
        )
    base_keys = [row[0] for row in LEGACY_FIELDS]
    extra_keys = [k for k in FIELDS if k not in base_keys and k != "id"]
    baseline_parts = [
        "roundtrip_baseline",
        *[f"roundtrip_baseline_{i}" for i in range(2, 33)],
    ]
    keys = [
        *base_keys,
        *extra_keys,
        "id",
        "record_revision",
        "window_opening_shots",
        "current_completed_shots",
        "current_good_units",
        "window_opening_good_units",
        "shift_quantity_unit",
        "opening_cutoff_at",
        "run_id",
        *baseline_parts,
    ]
    shift_keys = [
        f"{(start + timedelta(days=i)).isoformat()}|{s}"
        for i in range(payload.window_days)
        for s in ("DAY", "NIGHT")
    ]
    columns = [*keys, *shift_keys]
    settings = db.get(FactorySettings, payload.factory_id)
    as_of = now()
    metadata = {
        "schema_version": SCHEMA,
        "export_id": uuid4().hex,
        "factory_id": payload.factory_id,
        "business_timezone": "Asia/Shanghai",
        "source_revision": settings.revision,
        "shared_revision": revision(db, payload.factory_id)["shared_revision"],
        "as_of": as_of.isoformat(),
        "formula_version": "injection-v3-2",
        "date_window": {"start": start.isoformat(), "end_exclusive": end.isoformat()},
        "opening_cutoff_at": f"{start.isoformat()}T{settings.parameters['day_start']}:00+08:00",
        "quantity_cutoff_at": f"{end.isoformat()}T{settings.parameters['day_start']}:00+08:00",
        "field_keys": columns,
        "baseline_parts": baseline_parts,
        "formula_mode": payload.editable_formulas,
        "time_columns_note": "排程时间为服务器快照；修改数量后需回系统重算。",
    }
    links_by_demand = {}
    actual_run_ids = set()
    for link, run in db.execute(
        select(RunDemand, Run)
        .join(Run, Run.id == RunDemand.run_id)
        .where(RunDemand.demand_id.in_(ids), Run.actual_start_at.is_not(None))
    ):
        links_by_demand.setdefault(link.demand_id, []).append(
            (link, run, timestamp(run.actual_start_at), timestamp(run.actual_end_at))
        )
        actual_run_ids.add(run.id)
    run_counts = (
        dict(
            db.execute(
                select(RunDemand.run_id, func.count())
                .where(RunDemand.run_id.in_(actual_run_ids))
                .group_by(RunDemand.run_id)
            ).all()
        )
        if actual_run_ids
        else {}
    )
    from .reports import shift_window

    windows = []
    if actual_run_ids:
        for key in shift_keys:
            day, shift = key.split("|")
            start_at, end_at = shift_window(day, shift, settings.parameters)
            if start_at <= as_of:
                windows.append((key, day, shift, start_at, end_at))

    def targets(demand_id):
        links = links_by_demand.get(demand_id)
        if not links:
            return {}
        result = {}
        for key, day, shift, start_at, end_at in windows:
            if key in historical_keys[demand_id]:
                continue
            candidates = [
                (link, run)
                for link, run, actual_start, actual_end in links
                if actual_start < end_at and (not actual_end or actual_end > start_at)
            ]
            reports = report_metadata[demand_id].get(key, [])
            if len(candidates) != 1 or len(reports) > 1:
                continue
            link, run = candidates[0]
            if (
                run_counts.get(run.id) != 1
                or link.allocation_mode != "SEQUENTIAL_SHOTS"
            ):
                continue
            report = reports[0] if reports else None
            if report and (
                report["unit"] != "SHOT"
                or report["allocated"] != report["physical_shots"]
            ):
                continue
            result[key] = {
                "report_key": {
                    "run_id": run.id,
                    "production_date": day,
                    "shift_code": shift,
                    "segment_key": report["segment_key"] if report else "default",
                },
                "id": report["id"] if report else None,
                "revision": report["revision"] if report else 0,
            }
        return result

    wb = Workbook()
    plan = wb.active
    plan.title = "计划表"
    exchange = wb.create_sheet("标准交换表")
    header_fill = PatternFill("solid", fgColor="006559")
    machine_fill = PatternFill("solid", fgColor="E2F3F0")
    body_alignment = Alignment(vertical="center", wrap_text=False)
    header_alignment = Alignment(wrap_text=True, vertical="center")
    header_font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF", size=10)
    title_font = Font(name="Microsoft YaHei", size=16, bold=True)
    machines = {
        m.code: m
        for m in db.scalars(
            select(Machine).where(Machine.factory_id == payload.factory_id)
        )
    }

    def write_cell(ws, row, col, value, typ=None):
        cell = ws.cell(row, col)
        if isinstance(value, (dict, list)):
            value = json.dumps(
                jsonable(value), ensure_ascii=False, separators=(",", ":")
            )
        if typ == "datetime" and value:
            value = timestamp(value).replace(tzinfo=None)
            cell.number_format = "yyyy-mm-dd hh:mm"
        cell.value = value
        if isinstance(value, str):
            cell.data_type = "s"
        cell.alignment = body_alignment
        return cell

    for ws in (plan, exchange):
        ws.cell(
            1,
            1,
            f"{FACTORIES[payload.factory_id]}啤机生产日计划表"
            + (" · 标准交换" if ws is exchange else ""),
        )
        ws.cell(1, 1).font = title_font
        ws.cell(
            2,
            1,
            "INJECTION_V3:"
            + json.dumps(metadata, ensure_ascii=False, separators=(",", ":")),
        )
        ws.row_dimensions[2].hidden = True
        for col, key in enumerate(columns, 1):
            label = (
                key
                if ws is exchange
                else FIELDS[key].label
                if key in FIELDS
                else "窗口前已啤"
                if key == "window_opening_shots"
                else key
            )
            if "|" in key and ws is plan:
                d, s = key.split("|")
                label = f"{d}\n{'白班' if s == 'DAY' else '夜班'}"
            cell = write_cell(ws, 3, col, label)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            ws.column_dimensions[get_column_letter(col)].width = (
                17 if key in {"mold_code", "order_no", "item_no"} else 14
            )
            if key in {"id", "record_revision", "run_id", *baseline_parts}:
                ws.column_dimensions[get_column_letter(col)].hidden = True
        ws.freeze_panes = "H4"
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[3].height = 32
    plan_row = 4
    exchange_row = 4
    previous = None
    editable_keys = [key for key, spec in FIELDS.items() if spec.editable]
    column_types = [
        (col, key, FIELDS[key].value_type if key in FIELDS else None)
        for col, key in enumerate(columns, 1)
    ]
    opening_col = get_column_letter(columns.index("window_opening_shots") + 1)
    adjustment_col = get_column_letter(columns.index("adjustment_shots") + 1)
    allowance_col = get_column_letter(columns.index("allowance_rate") + 1)
    first_shift = get_column_letter(len(keys) + 1)
    last_shift = get_column_letter(len(columns))
    for demand in rows:
        values = record(demand)
        code = values.get("machine_code") or "待排需求"
        if code != previous:
            machine = machines.get(code)
            write_cell(plan, plan_row, 2, code)
            write_cell(
                plan,
                plan_row,
                7,
                f"{machine.machine_a}A" if machine and machine.machine_a else "待排",
            )
            write_cell(
                plan, plan_row, 8, machine.machine_family if machine else "需求池"
            )
            for col in range(1, min(50, len(columns)) + 1):
                plan.cell(plan_row, col).fill = machine_fill
            plan_row += 1
            previous = code
        window_total = sum(shifts[demand.id].values())
        current_completed = values.get("completed_shots")
        current_good = values.get("good_units")
        # A past-window export is explicitly a quantity snapshot at its cutoff;
        # later reports are never mislabeled as production before this window.
        if demand.allocation_mode != "CO_OUTPUT_UNITS":
            values.update(
                jsonable(
                    quantities(
                        {**values, "opening_shots": before[demand.id]},
                        decimal(window_total),
                    )
                )
            )
        else:
            values.update(
                jsonable(
                    quantities(
                        values, good_units=before_units[demand.id] + window_total
                    )
                )
            )
        values.update(
            id=demand.id,
            record_revision=demand.revision,
            run_id=values.get("active_run_id"),
            window_opening_shots=before[demand.id],
            current_completed_shots=current_completed,
            current_good_units=current_good,
            window_opening_good_units=before_units[demand.id],
            shift_quantity_unit="UNIT"
            if demand.allocation_mode == "CO_OUTPUT_UNITS"
            else "SHOT",
            opening_cutoff_at=metadata["opening_cutoff_at"],
            roundtrip_baseline={
                "fields": {key: values.get(key) for key in editable_keys},
                "reports": report_metadata[demand.id],
                "report_targets": targets(demand.id),
                "shift_values": shifts[demand.id],
            },
        )
        encoded = json.dumps(
            jsonable(values["roundtrip_baseline"]),
            ensure_ascii=False,
            separators=(",", ":"),
        )
        if len(encoded) > 30000 * len(baseline_parts):
            raise ValueError("报工分段元数据过大，请缩短导出日期窗口")
        for index, key in enumerate(baseline_parts):
            # Excel truncates cell strings at 32767 characters. Chunking keeps
            # long-window report revisions intact instead of invalid JSON.
            values[key] = encoded[index * 30000 : (index + 1) * 30000] or None
        values.update(shifts[demand.id])
        serialized = []
        for col, key, typ in column_types:
            value = values.get(key)
            if value is None:
                continue
            if isinstance(value, (dict, list)):
                value = json.dumps(
                    jsonable(value), ensure_ascii=False, separators=(",", ":")
                )
            if typ == "datetime" and value:
                value = timestamp(value).replace(tzinfo=None)
            serialized.append((col, value, typ))
        for ws, row_index in ((plan, plan_row), (exchange, exchange_row)):
            for col, value, typ in serialized:
                cell = ws.cell(row_index, col, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                if typ == "datetime" and value:
                    cell.number_format = "yyyy-mm-dd hh:mm"
                cell.alignment = body_alignment
            ws.row_dimensions[row_index].height = 24
            if (
                payload.editable_formulas
                and demand.allocation_mode != "CO_OUTPUT_UNITS"
            ):
                for col, formula in {
                    "M": f"={opening_col}{row_index}+SUM({first_shift}{row_index}:{last_shift}{row_index})",
                    "N": f'=IF(L{row_index}="","",MAX(0,L{row_index}+{adjustment_col}{row_index}-M{row_index}))',
                    "V": f'=IF(OR(N{row_index}="",T{row_index}=""),"",N{row_index}*T{row_index}/1000*(1+{allowance_col}{row_index}))',
                    "X": f'=IF(OR(N{row_index}="",W{row_index}=""),"",N{row_index}*W{row_index})',
                }.items():
                    ws[f"{col}{row_index}"] = formula
        plan_row += 1
        exchange_row += 1
    exchange.auto_filter.ref = (
        f"A3:{get_column_letter(len(columns))}{max(3, exchange_row - 1)}"
    )
    note = wb.create_sheet("口径说明")
    for item in [
        "后端数值为导出时的权威快照。",
        metadata["time_columns_note"],
        "SHOT 行：窗口前已啤 + 本窗口明细 = 已啤数；空白为未报，0 为已报零。",
        "UNIT 行：班次格为良品件，窗口前良品件 + 本窗口明细 = 累计良品件，不计入啤数 M；当前良品累计另见 current_good_units。",
        "数量按所选窗口截止日汇总；窗口后的报工不计入本表 M，当前累计另见 current_completed_shots。",
        "单需求且唯一物理批次的班次格可回传更正；组批分配格与导入历史格只读，请在系统班次报工面维护物理数量。",
        "历史单元格数量不等同已去重的机台物理啤数。同啤多产物见分配口径与良品件字段。",
        "公式版需要 Excel/WPS 重算，导出文件尚无新公式缓存。"
        if payload.editable_formulas
        else "本文件使用服务器数值，不包含 NOW、WEEKNUM 或外部工作簿引用。",
    ]:
        note.append([item])
    note.column_dimensions["A"].width = 110
    output = io.BytesIO()
    wb.save(output)
    return (
        output.getvalue(),
        f"{FACTORIES[payload.factory_id]}_注塑计划_{start.isoformat()}.xlsx",
    )
