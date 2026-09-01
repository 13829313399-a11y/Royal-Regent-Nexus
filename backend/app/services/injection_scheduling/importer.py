from __future__ import annotations

import hashlib
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from app.services.injection_scheduling.calculations import (
    estimated_material_kg,
    normalize_ratio,
    parse_machine_a_value,
)
from app.services.injection_scheduling.import_profiles import (
    IMPORT_PROFILES,
    ImportProfile,
    canonical_header,
)
from app.services.injection_scheduling.template_contract import (
    UNIFIED_IMPORT_PROFILE,
    UNIFIED_PLAN_HEADERS,
    UNIFIED_PLAN_SHEET,
    UNIFIED_REQUIRED_SHEETS,
    UNIFIED_TEMPLATE_TITLE,
    UNIFIED_TEMPLATE_VERSION,
    UNIFIED_TEMPLATE_VERSION_MARKER,
)

MAX_IMPORT_BYTES = 35 * 1024 * 1024
MAX_IMPORT_MEGABYTES = 35
MAX_PREVIEW_ROWS = 2_000
MAX_SCAN_COLUMNS = 600
EXCEL_ERROR_VALUES = {
    "#N/A",
    "#VALUE!",
    "#NAME?",
    "#REF!",
    "#DIV/0!",
    "#NUM!",
    "#NULL!",
}


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _compact(value: object) -> str:
    return re.sub(r"\s+", "", _text(value))


def _decimal(value: object) -> Decimal | None:
    text = _text(value).replace(",", "")
    if not text or text.upper() in EXCEL_ERROR_VALUES or text.startswith("="):
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _date_text(value: object, *, epoch: datetime) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and 20_000 <= float(value) <= 80_000:
        return from_excel(value, epoch).date().isoformat()
    text = _text(value)
    if not text or text.upper() in EXCEL_ERROR_VALUES or text.startswith("="):
        return ""
    normalized = text.replace("/", "-").replace(".", "-")
    for candidate in (normalized, normalized[:10]):
        try:
            return date.fromisoformat(candidate).isoformat()
        except ValueError:
            continue
    match = re.search(r"(20\d{2})[-年](\d{1,2})[-月](\d{1,2})", normalized)
    if match:
        return date(*(int(part) for part in match.groups())).isoformat()
    return ""


def _effective_values(
    raw_values: tuple[Any, ...], cached_values: tuple[Any, ...]
) -> tuple[Any, ...]:
    effective: list[Any] = []
    for index, raw_value in enumerate(raw_values):
        cached_value = cached_values[index] if index < len(cached_values) else None
        raw_text = _text(raw_value)
        use_cached = cached_value not in (None, "") and (
            raw_value in (None, "")
            or raw_text.startswith("=")
            or raw_text.upper() in EXCEL_ERROR_VALUES
        )
        effective.append(cached_value if use_cached else raw_value)
    return tuple(effective)


def _shift_code(value: object) -> str:
    normalized = _compact(value)
    if normalized == "白班":
        return "DAY"
    if normalized in {"夜班", "晚班"}:
        return "NIGHT"
    return ""


def _history_columns(
    sheet, *, header_row: int, epoch: datetime
) -> tuple[list[tuple[int, str, str]], bool]:
    columns: list[tuple[int, str, str]] = []
    unresolved = False
    previous_date = ""
    for column in range(1, min(sheet.max_column, MAX_SCAN_COLUMNS) + 1):
        upper_value = sheet.cell(max(header_row - 1, 1), column).value
        header_value = sheet.cell(header_row, column).value
        shift = _shift_code(header_value) or _shift_code(upper_value)
        if not shift:
            continue
        date_candidate = upper_value if _shift_code(header_value) else header_value
        production_date = _date_text(date_candidate, epoch=epoch)
        if shift == "DAY":
            previous_date = production_date
        elif not production_date:
            production_date = previous_date
        if not production_date:
            unresolved = True
            continue
        columns.append((column, production_date, shift))
    return columns, unresolved


def _issue(
    *,
    row: int | None,
    field: str,
    raw_value: object,
    severity: str,
    error_type: str,
    message: str,
    blocking: bool,
) -> dict[str, object]:
    return {
        "source_row": row,
        "field_name": field,
        "raw_value": _text(raw_value),
        "severity": severity,
        "error_type": error_type,
        "message": message,
        "blocking": blocking,
    }


def _detect_profile(workbook, override: str) -> ImportProfile:
    if override:
        profile = IMPORT_PROFILES.get(override)
        if profile is None:
            raise HTTPException(status_code=422, detail="不支持的导入 profile")
        return profile

    if UNIFIED_REQUIRED_SHEETS <= set(workbook.sheetnames):
        unified_sheet = workbook[UNIFIED_PLAN_SHEET]
        if _text(unified_sheet["A1"].value) == UNIFIED_TEMPLATE_TITLE:
            return IMPORT_PROFILES[UNIFIED_IMPORT_PROFILE]

    sheet = workbook[workbook.sheetnames[0]]
    sample_values = [
        _compact(value)
        for row in sheet.iter_rows(
            min_row=1,
            max_row=min(sheet.max_row, 30),
            max_col=min(sheet.max_column, 80),
            values_only=True,
        )
        for value in row
        if value is not None
    ]
    joined = "\n".join(sample_values)
    matches = [
        profile
        for profile in IMPORT_PROFILES.values()
        if profile.title_signals
        and any(_compact(signal) in joined for signal in profile.title_signals)
    ]
    if len(matches) == 1:
        return matches[0]

    header_values = set(sample_values)
    flat = IMPORT_PROFILES["WAREHOUSE_ORDER_FLAT_V1"]
    if all(_compact(header) in header_values for header in flat.required_headers):
        return flat
    raise HTTPException(
        status_code=422,
        detail="无法唯一识别 Excel 模板，请选择明确的导入 profile",
    )


def _profile_sheet(workbook, profile: ImportProfile):
    if profile.code == UNIFIED_IMPORT_PROFILE:
        if UNIFIED_PLAN_SHEET not in workbook.sheetnames:
            raise HTTPException(status_code=422, detail="统一模板缺少统一计划表")
        return workbook[UNIFIED_PLAN_SHEET]
    return workbook[workbook.sheetnames[0]]


def _validate_unified_workbook(workbook, factory_id: str) -> None:
    missing_sheets = sorted(UNIFIED_REQUIRED_SHEETS - set(workbook.sheetnames))
    if missing_sheets:
        raise HTTPException(
            status_code=422,
            detail=f"统一模板缺少工作表：{', '.join(missing_sheets)}",
        )
    if workbook.sheetnames != [UNIFIED_PLAN_SHEET]:
        raise HTTPException(
            status_code=422,
            detail="统一模板只能包含“统一计划表”一个工作表",
        )
    plan = workbook[UNIFIED_PLAN_SHEET]
    if _text(plan["A1"].value) != UNIFIED_TEMPLATE_TITLE:
        raise HTTPException(status_code=422, detail="统一计划表标题已改变")
    version_marker = _text(plan["A2"].value)
    if version_marker != UNIFIED_TEMPLATE_VERSION_MARKER:
        raise HTTPException(
            status_code=422,
            detail=(
                f"统一模板版本应为 {UNIFIED_TEMPLATE_VERSION}，"
                f"当前为 {version_marker or '空'}"
            ),
        )
    headers = tuple(
        _text(plan.cell(4, column).value)
        for column in range(1, len(UNIFIED_PLAN_HEADERS) + 1)
    )
    if headers != UNIFIED_PLAN_HEADERS:
        raise HTTPException(
            status_code=422,
            detail="统一计划表 56 字段或字段顺序已改变，请从模块重新下载模板",
        )


def _header_mapping(
    sheet, profile: ImportProfile
) -> tuple[int, dict[str, int], dict[str, str]]:
    required = {_compact(header) for header in profile.required_headers}
    best: tuple[int, dict[str, int], dict[str, str], int] | None = None
    for row_number, values in enumerate(
        sheet.iter_rows(
            min_row=1,
            max_row=min(sheet.max_row, 40),
            max_col=min(sheet.max_column, MAX_SCAN_COLUMNS),
            values_only=True,
        ),
        start=1,
    ):
        mapping: dict[str, int] = {}
        labels: dict[str, str] = {}
        raw_headers: set[str] = set()
        for column, raw_value in enumerate(values, start=1):
            value = _text(raw_value)
            if not value:
                continue
            raw_headers.add(_compact(value))
            field_name = canonical_header(value)
            if field_name and field_name not in mapping:
                mapping[field_name] = column
                labels[field_name] = value
        matched = len(required & raw_headers)
        if best is None or matched > best[3]:
            best = (row_number, mapping, labels, matched)
    if best is None or best[3] < max(2, len(required) - 1):
        raise HTTPException(status_code=422, detail="未找到可识别的明细标题行")
    return best[0], best[1], best[2]


def _cell(row: tuple[Any, ...], mapping: dict[str, int], field_name: str) -> object:
    column = mapping.get(field_name)
    return row[column - 1] if column and column <= len(row) else None


def _extract_labeled_values(sheet, labels: dict[str, str]) -> dict[str, str]:
    results = {field_name: "" for field_name in labels}
    normalized_labels = {
        field_name: _compact(label) for field_name, label in labels.items()
    }
    for row in sheet.iter_rows(
        min_row=1, max_row=min(sheet.max_row, 25), max_col=min(sheet.max_column, 40)
    ):
        for index, cell in enumerate(row):
            value = _text(cell.value)
            normalized = _compact(value)
            for field_name, normalized_label in normalized_labels.items():
                if results[field_name] or not normalized.startswith(normalized_label):
                    continue
                for separator in (":", "："):
                    if separator in value:
                        inline = value.split(separator, 1)[1].strip()
                        if inline:
                            results[field_name] = inline
                            break
                if results[field_name]:
                    continue
                for following in row[index + 1 : index + 4]:
                    if _text(following.value):
                        results[field_name] = _text(following.value)
                        break
    return results


def _business_key(factory_id: str, row: dict[str, object]) -> str:
    canonical = "\x1f".join(
        str(row.get(field) or "").strip()
        for field in (
            "factory_id",
            "order_no",
            "demand_line_no",
            "product_code",
            "mold_code",
            "color",
            "material_name",
        )
    )
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def _normalize_row(
    *,
    factory_id: str,
    sheet_name: str,
    source_row: int,
    values: tuple[Any, ...],
    raw_values: tuple[Any, ...],
    mapping: dict[str, int],
    epoch: datetime,
    header: dict[str, str],
    machine_context: dict[str, object] | None,
) -> tuple[dict[str, object], list[dict[str, object]]]:
    raw = {field: _cell(values, mapping, field) for field in mapping}
    source_raw = {field: _cell(raw_values, mapping, field) for field in mapping}
    row_issues: list[dict[str, object]] = []
    water_ratio = None
    try:
        water_ratio = normalize_ratio(raw.get("water_ratio"))
    except ValueError:
        row_issues.append(
            _issue(
                row=source_row,
                field="water_ratio",
                raw_value=raw.get("water_ratio"),
                severity="WARNING",
                error_type="INVALID_RATIO",
                message="水口比例无法解析，将使用厂区默认值",
                blocking=False,
            )
        )

    quantity_sets = _decimal(raw.get("quantity_sets"))
    total_sets = _decimal(raw.get("total_sets"))
    order_shots = _decimal(raw.get("order_shots"))
    if order_shots is None and quantity_sets is not None:
        order_shots = quantity_sets
        row_issues.append(
            _issue(
                row=source_row,
                field="order_shots",
                raw_value=raw.get("order_shots"),
                severity="WARNING",
                error_type="DERIVED_ORDER_SHOTS",
                message="来源未给出啤数，预览暂以数量代替，确认前需核对",
                blocking=False,
            )
        )

    required_label = _text(raw.get("required_machine_a_label"))
    if not required_label and machine_context:
        required_label = _text(machine_context.get("machine_a_label"))
        if required_label:
            row_issues.append(
                _issue(
                    row=source_row,
                    field="required_machine_a_value",
                    raw_value=required_label,
                    severity="WARNING",
                    error_type="INFERRED_FROM_LEGACY_MACHINE_BLOCK",
                    message="来源缺少模具机安，暂按原计划所在机台机安迁移，需在模具主数据复核",
                    blocking=False,
                )
            )
    required_value = parse_machine_a_value(required_label)
    if required_label.upper() in EXCEL_ERROR_VALUES or required_label.startswith("="):
        required_label = ""
        required_value = None

    delivery_due = _date_text(
        raw.get("delivery_due_date") or header.get("delivery_due_date"), epoch=epoch
    )
    remark = _text(raw.get("remark"))
    raw_priority = _compact(raw.get("priority")).upper()
    priority = (
        "EXPEDITE"
        if raw_priority in {"特急", "EXPEDITE"} or "特急" in remark
        else "URGENT"
        if raw_priority in {"急", "急单", "URGENT"} or "急" in remark
        else "NORMAL"
    )
    raw_lightness = _compact(raw.get("color_lightness")).upper()
    color_lightness = {
        "透明": "TRANSPARENT",
        "浅": "LIGHT",
        "中": "MEDIUM",
        "深": "DARK",
        "未知": "UNKNOWN",
        "TRANSPARENT": "TRANSPARENT",
        "LIGHT": "LIGHT",
        "MEDIUM": "MEDIUM",
        "DARK": "DARK",
        "UNKNOWN": "UNKNOWN",
    }.get(raw_lightness, "UNKNOWN")
    raw_material_status = _compact(raw.get("material_status")).upper()
    material_status = {
        "未配料": "UNPREPARED",
        "部分配料": "PARTIAL",
        "已齐料": "READY",
        "缺料": "BLOCKED",
        "UNPREPARED": "UNPREPARED",
        "PARTIAL": "PARTIAL",
        "READY": "READY",
        "BLOCKED": "BLOCKED",
    }.get(raw_material_status, "UNPREPARED")
    product_code = _text(raw.get("product_code"))
    mold_code = _text(raw.get("mold_code"))
    order_no = _text(raw.get("order_no") or header.get("order_no"))
    material_weight = _decimal(raw.get("material_weight_kg"))
    if material_weight is None:
        material_weight = estimated_material_kg(
            order_shots,
            _decimal(raw.get("net_weight_g")),
            water_ratio,
        )

    normalized: dict[str, object] = {
        "factory_id": factory_id,
        "source_sheet": sheet_name,
        "source_row": source_row,
        "demand_line_no": str(source_row),
        "machine_code": _text(raw.get("machine_code"))
        or _text((machine_context or {}).get("machine_code")),
        "order_no": order_no,
        "product_code": product_code,
        "mold_code": mold_code,
        "product_name": _text(raw.get("product_name")),
        "quantity_sets": str(quantity_sets or Decimal(0)),
        "total_sets": str(total_sets) if total_sets is not None else None,
        "order_shots": str(order_shots or Decimal(0)),
        "qualified_shots": str(_decimal(raw.get("qualified_shots")) or Decimal(0)),
        "priority": priority,
        "order_date": _date_text(
            raw.get("order_date") or header.get("order_date"), epoch=epoch
        ),
        "delivery_start_date": _date_text(raw.get("delivery_start_date"), epoch=epoch),
        "delivery_due_date": delivery_due,
        "shipping_date": _date_text(raw.get("shipping_date"), epoch=epoch),
        "warehouse": _text(raw.get("warehouse") or header.get("warehouse")),
        "delivery_location": header.get("delivery_location", ""),
        "ordered_by_name": header.get("ordered_by_name", ""),
        "operator_name": header.get("operator_name", ""),
        "color": _text(raw.get("color")),
        "pigment_code": _text(raw.get("pigment_code")),
        "color_lightness": color_lightness,
        "material_name": _text(raw.get("material_name")),
        "water_ratio": str(water_ratio) if water_ratio is not None else None,
        "net_weight_g": str(_decimal(raw.get("net_weight_g")))
        if _decimal(raw.get("net_weight_g")) is not None
        else None,
        "gross_weight_g": str(_decimal(raw.get("gross_weight_g")))
        if _decimal(raw.get("gross_weight_g")) is not None
        else None,
        "material_weight_kg": str(material_weight)
        if material_weight is not None
        else None,
        "daily_target": str(_decimal(raw.get("daily_target")))
        if _decimal(raw.get("daily_target")) is not None
        else None,
        "material_status": material_status,
        "material_prepared_kg": str(
            _decimal(raw.get("material_prepared_kg")) or Decimal(0)
        ),
        "required_machine_a_label": required_label,
        "required_machine_a_value": str(required_value)
        if required_value is not None
        else None,
        "spray_required": "喷油" in remark,
        "remark": remark,
        "raw_source": {field: _text(value) for field, value in source_raw.items()},
    }

    required_checks = (
        ("order_no", order_no, "缺少单号"),
        ("product_code", product_code, "缺少货号/款号"),
        ("mold_code", mold_code, "缺少工模/模具编号"),
        ("order_shots", order_shots, "数量/啤数无法确定"),
        ("delivery_due_date", delivery_due, "交货期无法解析"),
        (
            "required_machine_a_value",
            required_value,
            "必需机安缺失且无法从当前文件补齐",
        ),
    )
    for field_name, value, message in required_checks:
        if value in (None, "", Decimal(0)):
            row_issues.append(
                _issue(
                    row=source_row,
                    field=field_name,
                    raw_value=normalized.get(field_name),
                    severity="ERROR",
                    error_type="REQUIRED_FIELD_MISSING",
                    message=message,
                    blocking=True,
                )
            )

    if water_ratio is None:
        row_issues.append(
            _issue(
                row=source_row,
                field="water_ratio",
                raw_value=raw.get("water_ratio"),
                severity="WARNING",
                error_type="DEFAULT_WATER_RATIO",
                message="水口比例缺失，确认导入时使用厂区默认值",
                blocking=False,
            )
        )
    normalized["data_completeness_status"] = (
        "BLOCKED" if any(item["blocking"] for item in row_issues) else "COMPLETE"
    )
    normalized["business_key"] = _business_key(factory_id, normalized)
    return normalized, row_issues


def _looks_like_machine_header(
    values: tuple[Any, ...], mapping: dict[str, int]
) -> bool:
    if (
        _decimal(_cell(values, mapping, "quantity_sets")) is not None
        or _decimal(_cell(values, mapping, "order_shots")) is not None
    ):
        return False
    first_values = [_text(value) for value in values[:14] if _text(value)]
    return any(parse_machine_a_value(value) is not None for value in first_values)


def parse_workbook_preview(
    *,
    content: bytes,
    file_name: str,
    factory_id: str,
    profile_override: str = "",
) -> dict[str, object]:
    if not content:
        raise HTTPException(status_code=400, detail="请选择有效的 Excel 文件")
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(
            status_code=413,
            detail=f"排产导入文件不能超过 {MAX_IMPORT_MEGABYTES} MB",
        )
    if Path(file_name).suffix.lower() not in {".xlsx", ".xlsm"}:
        raise HTTPException(status_code=400, detail="排产导入仅支持 .xlsx 或 .xlsm")

    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise HTTPException(
            status_code=400, detail="Excel 文件损坏或格式不受支持"
        ) from exc
    try:
        cached_workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:
        workbook.close()
        raise HTTPException(
            status_code=400, detail="Excel 文件损坏或格式不受支持"
        ) from exc

    try:
        profile = _detect_profile(workbook, profile_override)
        if profile.code == UNIFIED_IMPORT_PROFILE:
            _validate_unified_workbook(workbook, factory_id)
        sheet = _profile_sheet(workbook, profile)
        cached_sheet = cached_workbook[sheet.title]
        header_row, mapping, labels = _header_mapping(sheet, profile)
        history_columns, unresolved_history_dates = (
            _history_columns(sheet, header_row=header_row, epoch=workbook.epoch)
            if profile.document_type == "PLAN_MIGRATION"
            else ([], False)
        )
        common_header = _extract_labeled_values(
            sheet,
            {
                "order_no": "单号",
                "delivery_due_date": "交货日期",
                "delivery_location": "交货地点",
                "warehouse": "仓库",
                "ordered_by_name": "下单人",
                "operator_name": "操作员",
                "company_name": "公司名称",
            },
        )
        rows: list[dict[str, object]] = []
        issues: list[dict[str, object]] = []
        machines: dict[str, dict[str, object]] = {}
        history_outputs: list[dict[str, object]] = []
        current_machine: dict[str, object] | None = None
        max_row = min(sheet.max_row, header_row + MAX_PREVIEW_ROWS)
        raw_rows = sheet.iter_rows(
            min_row=header_row + 1,
            max_row=max_row,
            max_col=min(sheet.max_column, MAX_SCAN_COLUMNS),
            values_only=True,
        )
        cached_rows = cached_sheet.iter_rows(
            min_row=header_row + 1,
            max_row=max_row,
            max_col=min(cached_sheet.max_column, MAX_SCAN_COLUMNS),
            values_only=True,
        )
        for source_row, (raw_values, cached_values) in enumerate(
            zip(raw_rows, cached_rows, strict=False),
            start=header_row + 1,
        ):
            values = _effective_values(raw_values, cached_values)
            if not any(value not in (None, "") for value in values):
                continue
            if profile.document_type == "PLAN_MIGRATION" and _looks_like_machine_header(
                values, mapping
            ):
                machine_code = _text(_cell(values, mapping, "machine_code"))
                machine_label = next(
                    (
                        _text(value)
                        for value in values[:14]
                        if parse_machine_a_value(_text(value)) is not None
                    ),
                    "",
                )
                if machine_code:
                    current_machine = {
                        "machine_code": machine_code,
                        "machine_a_label": machine_label,
                        "machine_a_value": str(parse_machine_a_value(machine_label)),
                        "source_row": source_row,
                    }
                    machines[machine_code] = current_machine
                continue

            product_code = _text(_cell(values, mapping, "product_code"))
            mold_code = _text(_cell(values, mapping, "mold_code"))
            if not any((product_code, mold_code)):
                continue
            if profile.code == UNIFIED_IMPORT_PROFILE:
                row_factory_id = _text(_cell(values, mapping, "factory_id")).lower()
                if row_factory_id and row_factory_id != factory_id:
                    issues.append(
                        _issue(
                            row=source_row,
                            field="factory_id",
                            raw_value=row_factory_id,
                            severity="ERROR",
                            error_type="FACTORY_SCOPE_MISMATCH",
                            message="明细工厂 ID 与当前页面厂区不一致",
                            blocking=True,
                        )
                    )
            if (
                profile.document_type == "ORDER_FORM"
                and _decimal(_cell(values, mapping, "quantity_sets")) is None
                and _decimal(_cell(values, mapping, "order_shots")) is None
                and not _text(_cell(values, mapping, "product_name"))
            ):
                # The supplied form has summary/footer rows under the same merged
                # columns as detail rows. A detail row must retain at least one
                # quantity signal or an explicit product name.
                continue
            normalized, row_issues = _normalize_row(
                factory_id=factory_id,
                sheet_name=sheet.title,
                source_row=source_row,
                values=values,
                raw_values=raw_values,
                mapping=mapping,
                epoch=workbook.epoch,
                header=common_header,
                machine_context=current_machine,
            )
            rows.append(normalized)
            issues.extend(row_issues)
            for column, production_date, shift in history_columns:
                reported_shots = _decimal(
                    values[column - 1] if column <= len(values) else None
                )
                if reported_shots is None or reported_shots == 0:
                    continue
                if reported_shots < 0:
                    issues.append(
                        _issue(
                            row=source_row,
                            field=f"shift_output:{production_date}:{shift}",
                            raw_value=reported_shots,
                            severity="ERROR",
                            error_type="NEGATIVE_SHIFT_OUTPUT",
                            message="白夜班历史啤数不能为负数",
                            blocking=True,
                        )
                    )
                    continue
                history_outputs.append(
                    {
                        "business_key": normalized["business_key"],
                        "source_row": source_row,
                        "production_date": production_date,
                        "shift": shift,
                        "reported_shots": str(reported_shots),
                    }
                )

        if unresolved_history_dates:
            issues.append(
                _issue(
                    row=header_row,
                    field="shift_history_date",
                    raw_value="",
                    severity="WARNING",
                    error_type="HISTORY_DATE_UNRESOLVED",
                    message=(
                        "部分白夜班列只有日号、没有可确认的年月，预览未迁移这些列；"
                        "请使用含完整日期的历史模板"
                    ),
                    blocking=False,
                )
            )

        if not rows:
            raise HTTPException(
                status_code=422, detail="识别到模板，但未找到可预览的订单明细"
            )
        duplicate_keys = {
            key
            for key in {row["business_key"] for row in rows}
            if sum(item["business_key"] == key for item in rows) > 1
        }
        for key in duplicate_keys:
            duplicate_rows = [
                int(item["source_row"]) for item in rows if item["business_key"] == key
            ]
            issues.append(
                _issue(
                    row=min(duplicate_rows),
                    field="business_key",
                    raw_value=key,
                    severity="ERROR",
                    error_type="DUPLICATE_BUSINESS_KEY",
                    message=f"同一业务键在来源行 {duplicate_rows} 重复",
                    blocking=True,
                )
            )

        return {
            "profile_code": profile.code,
            "document_type": profile.document_type,
            "source_file_name": Path(file_name).name,
            "source_file_sha256": hashlib.sha256(content).hexdigest(),
            "source_size_bytes": len(content),
            "source_sheet": sheet.title,
            "header_row": header_row,
            "header": common_header,
            "column_mapping": labels,
            "machines": list(machines.values()),
            "rows": rows,
            "history_outputs": history_outputs,
            "issues": issues,
            "summary": {
                "row_count": len(rows),
                "machine_count": len(machines),
                "history_output_count": len(history_outputs),
                "blocking_issue_count": sum(bool(item["blocking"]) for item in issues),
                "warning_count": sum(item["severity"] == "WARNING" for item in issues),
            },
        }
    finally:
        workbook.close()
        cached_workbook.close()
