from __future__ import annotations

import math
import os
import re
import subprocess
import unicodedata
from csv import DictReader
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from statistics import median
from typing import Iterable

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


MAX_PDF_PAGES = 80
PROJECT_ROOT = Path(__file__).resolve().parents[3]
PDF_TESSDATA_DIR = PROJECT_ROOT / "tools" / "Tesseract-OCR" / "tessdata"
LAYOUT_COLUMN_COUNT = 36
LAYOUT_ROW_COUNT = 84
_TABLE_SETTINGS = {
    "vertical_strategy": "lines",
    "horizontal_strategy": "lines",
    "snap_tolerance": 4,
    "join_tolerance": 4,
    "intersection_tolerance": 5,
}
_TEXT_TABLE_SETTINGS = {
    "vertical_strategy": "text",
    "horizontal_strategy": "text",
    "min_words_vertical": 2,
    "min_words_horizontal": 1,
    "text_x_tolerance": 3,
    "text_y_tolerance": 3,
}


class PdfToExcelConversionError(ValueError):
    pass


@dataclass(frozen=True)
class LayoutBlock:
    text: str
    x0: float
    top: float
    x1: float
    bottom: float
    font_size: float = 9
    bold: bool = False


@dataclass(frozen=True)
class ExtractedSheet:
    title: str
    rows: list[list[str]]
    source: str
    layout_blocks: tuple[LayoutBlock, ...] = ()
    page_width: float = 0
    page_height: float = 0


@dataclass(frozen=True)
class PdfToExcelResult:
    content: bytes
    page_count: int
    table_count: int
    text_page_count: int
    ocr_page_count: int
    output_file_name: str


def _clean_cell(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\x00", " ").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
    return "\n".join(line for line in lines if line)


def _normalize_rows(rows: Iterable[Iterable[object]]) -> list[list[str]]:
    normalized = [[_clean_cell(cell) for cell in row] for row in rows]
    normalized = [row for row in normalized if any(cell for cell in row)]
    if not normalized:
        return []

    last_used_column = max(
        (index for row in normalized for index, cell in enumerate(row) if cell),
        default=0,
    )
    return [row[: last_used_column + 1] + [""] * max(0, last_used_column + 1 - len(row)) for row in normalized]


def _is_cjk_character(value: str) -> bool:
    return bool(value) and (
        "\u3400" <= value <= "\u4dbf"
        or "\u4e00" <= value <= "\u9fff"
        or "\uf900" <= value <= "\ufaff"
    )


def _join_tokens(tokens: Iterable[str]) -> str:
    text = ""
    for raw_token in tokens:
        token = _clean_cell(raw_token)
        if not token:
            continue
        if not text:
            text = token
            continue
        if token[0] in ",.;:!?%)]}>，。：；、）》】" or text[-1] in "([{</《【":
            text += token
        elif _is_cjk_character(text[-1]) and _is_cjk_character(token[0]):
            text += token
        else:
            text += f" {token}"
    return text


def _is_plausible_text_table(rows: list[list[str]]) -> bool:
    if len(rows) < 2:
        return False

    column_count = max((len(row) for row in rows), default=0)
    if column_count < 2:
        return False

    populated_per_row = [sum(bool(cell) for cell in row) for row in rows]
    if len(rows) >= 8 and median(populated_per_row) < 2:
        return False

    populated_cells = [cell for row in rows for cell in row if cell]
    if len(populated_cells) >= 12:
        short_fragments = sum(len(re.sub(r"\s+", "", cell)) <= 2 for cell in populated_cells)
        if short_fragments / len(populated_cells) > 0.26:
            return False
        short_alpha_fragments = sum(
            bool(re.fullmatch(r"[A-Za-z]{1,3}[.:]?", cell.strip()))
            for cell in populated_cells
        )
        if short_alpha_fragments / len(populated_cells) > 0.2:
            return False

    recurring_columns = 0
    for column_index in range(column_count):
        occurrences = sum(
            column_index < len(row) and bool(row[column_index])
            for row in rows
        )
        if occurrences >= max(2, math.ceil(len(rows) * 0.35)):
            recurring_columns += 1
    return recurring_columns >= 2


def _native_text_is_unreliable(text: str) -> bool:
    meaningful = [character for character in text if not character.isspace()]
    if not meaningful:
        return False

    invalid_count = sum(
        character == "\ufffd"
        or unicodedata.category(character) in {"Co", "Cs"}
        or (unicodedata.category(character) == "Cc" and character not in "\n\r\t")
        for character in meaningful
    )
    if invalid_count / len(meaningful) > 0.01:
        return True

    cjk_characters = [character for character in meaningful if _is_cjk_character(character)]
    if len(cjk_characters) < 40:
        return False

    unique_ratio = len(set(cjk_characters)) / len(cjk_characters)
    repeated_ratio = sum(
        left == right for left, right in zip(cjk_characters, cjk_characters[1:])
    ) / max(1, len(cjk_characters) - 1)
    return unique_ratio < 0.2 and repeated_ratio > 0.08


def _layout_block_from_words(words: list[dict]) -> LayoutBlock | None:
    text = _join_tokens(str(word.get("text", "")) for word in words)
    if not text:
        return None

    font_sizes = []
    font_names = []
    for word in words:
        try:
            font_sizes.append(float(word.get("size") or 9))
        except (TypeError, ValueError):
            pass
        font_names.append(str(word.get("fontname") or ""))

    return LayoutBlock(
        text=text,
        x0=min(float(word["x0"]) for word in words),
        top=min(float(word["top"]) for word in words),
        x1=max(float(word["x1"]) for word in words),
        bottom=max(float(word["bottom"]) for word in words),
        font_size=max(font_sizes, default=9),
        bold=any("bold" in font_name.lower() for font_name in font_names),
    )


def _extract_native_layout(page) -> list[LayoutBlock]:
    if not hasattr(page, "extract_words"):
        return []

    try:
        words = page.extract_words(
            x_tolerance=2,
            y_tolerance=2,
            keep_blank_chars=False,
            use_text_flow=False,
            extra_attrs=["fontname", "size"],
        ) or []
    except TypeError:
        try:
            words = page.extract_words(x_tolerance=2, y_tolerance=2) or []
        except Exception:
            return []
    except Exception:
        return []

    usable_words = [
        word
        for word in words
        if _clean_cell(word.get("text", ""))
        and all(key in word for key in ("x0", "x1", "top", "bottom"))
    ]
    if not usable_words:
        return []

    lines: list[list[dict]] = []
    line_tops: list[float] = []
    for word in sorted(usable_words, key=lambda item: (float(item["top"]), float(item["x0"]))):
        word_top = float(word["top"])
        try:
            word_size = float(word.get("size") or 9)
        except (TypeError, ValueError):
            word_size = 9
        tolerance = max(2.5, min(4.5, word_size * 0.35))
        if lines and abs(word_top - line_tops[-1]) <= tolerance:
            lines[-1].append(word)
            line_tops[-1] = sum(float(item["top"]) for item in lines[-1]) / len(lines[-1])
        else:
            lines.append([word])
            line_tops.append(word_top)

    blocks: list[LayoutBlock] = []
    for line in lines:
        ordered = sorted(line, key=lambda item: float(item["x0"]))
        groups: list[list[dict]] = []
        for word in ordered:
            if not groups:
                groups.append([word])
                continue
            previous = groups[-1][-1]
            try:
                previous_size = float(previous.get("size") or 9)
                current_size = float(word.get("size") or 9)
            except (TypeError, ValueError):
                previous_size = current_size = 9
            gap = float(word["x0"]) - float(previous["x1"])
            join_gap = max(4, min(10, (previous_size + current_size) * 0.45))
            if gap <= join_gap:
                groups[-1].append(word)
            else:
                groups.append([word])

        for group in groups:
            if block := _layout_block_from_words(group):
                blocks.append(block)

    return blocks


def _rows_from_text(text: str) -> list[list[str]]:
    rows: list[list[str]] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        cells = [_clean_cell(cell) for cell in re.split(r"\t+|\s{2,}", line)]
        rows.append(cells if any(cells) else [line])
    return _normalize_rows(rows)


def _extract_native_page(page, page_number: int) -> tuple[list[ExtractedSheet], str]:
    text = page.extract_text(x_tolerance=3, y_tolerance=3) or ""
    try:
        tables = page.extract_tables(table_settings=_TABLE_SETTINGS) or []
    except Exception:
        tables = []

    normalized_tables = [_normalize_rows(table) for table in tables]
    normalized_tables = [table for table in normalized_tables if table]

    if not normalized_tables and text.strip():
        try:
            text_tables = page.extract_tables(table_settings=_TEXT_TABLE_SETTINGS) or []
        except Exception:
            text_tables = []
        normalized_tables = [
            normalized
            for table in text_tables
            if (normalized := _normalize_rows(table))
            and _is_plausible_text_table(normalized)
        ]

    sheets = [
        ExtractedSheet(
            title=f"第{page_number}页_表格{table_index}",
            rows=table,
            source="table",
        )
        for table_index, table in enumerate(normalized_tables, start=1)
    ]
    return sheets, text


def _pdf_tessdata_directory() -> Path | None:
    configured = os.getenv("PDF_TO_EXCEL_TESSDATA_DIR", "").strip()
    candidates = [Path(configured)] if configured else []
    candidates.append(PDF_TESSDATA_DIR)
    for candidate in candidates:
        if candidate.is_dir() and any(candidate.glob("*.traineddata")):
            return candidate
    return None


def _tesseract_ocr_options(pytesseract_module) -> tuple[str, str, str, Path | None]:
    from app.services.carton_mark import configure_tesseract

    tesseract_cmd, default_language = configure_tesseract(pytesseract_module)
    if not tesseract_cmd:
        return "", "", "", None

    tessdata_directory = _pdf_tessdata_directory()
    if tessdata_directory is None:
        return (
            tesseract_cmd,
            default_language,
            "--oem 3 --psm 6 -c preserve_interword_spaces=1",
            None,
        )

    available = {path.stem for path in tessdata_directory.glob("*.traineddata")}
    preferred = [language for language in ("chi_tra", "chi_sim", "eng") if language in available]
    if not preferred:
        return (
            tesseract_cmd,
            default_language,
            "--oem 3 --psm 6 -c preserve_interword_spaces=1",
            None,
        )

    language = "+".join(preferred)
    return (
        tesseract_cmd,
        language,
        "--oem 1 --psm 6 -c preserve_interword_spaces=1",
        tessdata_directory,
    )


def _run_bundled_tesseract(
    image,
    *,
    tesseract_cmd: str,
    language: str,
    tessdata_directory: Path,
    output_format: str,
    timeout: int,
) -> str:
    image_bytes = BytesIO()
    image.save(image_bytes, format="PNG")
    command = [
        tesseract_cmd,
        "stdin",
        "stdout",
        "-l",
        language,
        "--tessdata-dir",
        str(tessdata_directory),
        "--oem",
        "1",
        "--psm",
        "3",
        "-c",
        "preserve_interword_spaces=1",
    ]
    if output_format == "tsv":
        command.extend(["-c", "tessedit_create_tsv=1"])
    completed = subprocess.run(
        command,
        input=image_bytes.getvalue(),
        capture_output=True,
        timeout=timeout,
        check=False,
    )
    if completed.returncode != 0:
        return ""
    return completed.stdout.decode("utf-8", errors="replace")


def _tsv_to_data(tsv_text: str) -> dict[str, list[object]]:
    fields = (
        "text",
        "conf",
        "block_num",
        "par_num",
        "line_num",
        "left",
        "top",
        "width",
        "height",
    )
    data: dict[str, list[object]] = {field: [] for field in fields}
    for row in DictReader(StringIO(tsv_text), delimiter="\t"):
        for field in fields:
            data[field].append(row.get(field, ""))
    return data


def _ocr_page_layout(
    pdf_bytes: bytes,
    page_index: int,
    page_width: float,
    page_height: float,
) -> list[LayoutBlock]:
    if page_width <= 0 or page_height <= 0:
        return []

    try:
        import pypdfium2 as pdfium  # type: ignore
        import pytesseract  # type: ignore
        from pytesseract import Output  # type: ignore
    except Exception:
        return []

    tesseract_cmd, language, config, tessdata_directory = _tesseract_ocr_options(pytesseract)
    if not tesseract_cmd or not language:
        return []

    document = None
    try:
        document = pdfium.PdfDocument(pdf_bytes)
        image = document[page_index].render(scale=4).to_pil().convert("RGB")
        if tessdata_directory is not None:
            data = _tsv_to_data(_run_bundled_tesseract(
                image,
                tesseract_cmd=tesseract_cmd,
                language=language,
                tessdata_directory=tessdata_directory,
                output_format="tsv",
                timeout=60,
            ))
        else:
            data = pytesseract.image_to_data(
                image,
                lang=language,
                config=config,
                output_type=Output.DICT,
                timeout=60,
            )
    except TypeError:
        if "image" not in locals() or tessdata_directory is not None:
            return []
        try:
            data = pytesseract.image_to_data(
                image,
                lang=language,
                config=config,
                output_type=Output.DICT,
            )
        except Exception:
            return []
    except Exception:
        return []
    finally:
        if document is not None:
            document.close()

    grouped: dict[tuple[int, int, int], list[dict[str, object]]] = {}
    texts = data.get("text", [])
    for index, raw_text in enumerate(texts):
        text = _clean_cell(raw_text)
        if not text:
            continue
        try:
            confidence = float(data.get("conf", [])[index])
        except (IndexError, TypeError, ValueError):
            confidence = 0
        if confidence < 20:
            continue
        try:
            key = (
                int(data.get("block_num", [])[index]),
                int(data.get("par_num", [])[index]),
                int(data.get("line_num", [])[index]),
            )
            left = float(data.get("left", [])[index])
            top = float(data.get("top", [])[index])
            width = float(data.get("width", [])[index])
            height = float(data.get("height", [])[index])
        except (IndexError, TypeError, ValueError):
            continue
        grouped.setdefault(key, []).append({
            "text": text,
            "left": left,
            "top": top,
            "right": left + width,
            "bottom": top + height,
        })

    scale_x = page_width / max(1, image.width)
    scale_y = page_height / max(1, image.height)
    blocks: list[LayoutBlock] = []
    for words in grouped.values():
        ordered = sorted(words, key=lambda item: float(item["left"]))
        text = _join_tokens(str(word["text"]) for word in ordered)
        if not text:
            continue
        top = min(float(word["top"]) for word in ordered) * scale_y
        bottom = max(float(word["bottom"]) for word in ordered) * scale_y
        font_size = max(8, min(20, (bottom - top) * 0.9))
        x0 = min(float(word["left"]) for word in ordered) * scale_x
        x1 = max(float(word["right"]) for word in ordered) * scale_x
        if top < page_height * 0.09 and x0 > page_width * 0.65 and len(text) <= 12:
            continue
        blocks.append(LayoutBlock(
            text=text,
            x0=x0,
            top=top,
            x1=x1,
            bottom=bottom,
            font_size=font_size,
            bold=bool(re.match(
                r"^(Release Order|Remarks:|Seller:|Vendor:|Quantity|Main Mark|Side Mark|Port of discharge)",
                text,
                re.IGNORECASE,
            )),
        ))

    blocks.sort(key=lambda block: (block.top, block.x0))
    combined_text = "\n".join(block.text for block in blocks)
    if _native_text_is_unreliable(combined_text):
        return []
    return blocks


def _ocr_page(pdf_bytes: bytes, page_index: int) -> list[list[str]]:
    try:
        import pypdfium2 as pdfium  # type: ignore
        import pytesseract  # type: ignore
    except Exception:
        return []

    tesseract_cmd, language, config, tessdata_directory = _tesseract_ocr_options(pytesseract)
    if not tesseract_cmd or not language:
        return []

    document = None
    try:
        document = pdfium.PdfDocument(pdf_bytes)
        image = document[page_index].render(scale=3).to_pil().convert("RGB")
        if tessdata_directory is not None:
            text = _run_bundled_tesseract(
                image,
                tesseract_cmd=tesseract_cmd,
                language=language,
                tessdata_directory=tessdata_directory,
                output_format="text",
                timeout=45,
            )
        else:
            text = pytesseract.image_to_string(
                image,
                lang=language,
                config=config,
                timeout=45,
            )
        return _rows_from_text(text)
    except Exception:
        return []
    finally:
        if document is not None:
            document.close()


def _safe_sheet_title(title: str, existing: set[str]) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", "_", title).strip()[:31] or "数据"
    candidate = cleaned
    suffix = 2
    while candidate in existing:
        suffix_text = f"_{suffix}"
        candidate = f"{cleaned[: 31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    existing.add(candidate)
    return candidate


def _style_data_sheet(worksheet, rows: list[list[str]]) -> None:
    worksheet.freeze_panes = "A2" if len(rows) > 1 else None
    worksheet.sheet_view.showGridLines = False
    if rows:
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        worksheet.auto_filter.ref = worksheet.dimensions

    for row in worksheet.iter_rows():
        for cell in row:
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column_index in range(1, worksheet.max_column + 1):
        values = [str(worksheet.cell(row=row_index, column=column_index).value or "") for row_index in range(1, worksheet.max_row + 1)]
        width = min(max((max((len(line) for line in value.splitlines()), default=0) for value in values), default=8) + 2, 48)
        worksheet.column_dimensions[get_column_letter(column_index)].width = max(width, 10)


def _style_layout_sheet(worksheet, extracted: ExtractedSheet) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 80
    worksheet.freeze_panes = None
    worksheet.page_setup.orientation = "portrait"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 1
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25
    worksheet.page_margins.top = 0.3
    worksheet.page_margins.bottom = 0.3

    for column_index in range(1, LAYOUT_COLUMN_COUNT + 1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = 2.65

    base_row_height = max(9, extracted.page_height / LAYOUT_ROW_COUNT)
    for row_index in range(1, LAYOUT_ROW_COUNT + 1):
        worksheet.row_dimensions[row_index].height = base_row_height

    mapped: dict[int, list[tuple[int, int, LayoutBlock]]] = {}
    for block in extracted.layout_blocks:
        row_index = max(
            1,
            min(
                LAYOUT_ROW_COUNT,
                round(block.top / max(1, extracted.page_height) * (LAYOUT_ROW_COUNT - 1)) + 1,
            ),
        )
        start_column = max(
            1,
            min(
                LAYOUT_COLUMN_COUNT,
                math.floor(block.x0 / max(1, extracted.page_width) * LAYOUT_COLUMN_COUNT) + 1,
            ),
        )
        end_column = max(
            start_column,
            min(
                LAYOUT_COLUMN_COUNT,
                math.ceil(block.x1 / max(1, extracted.page_width) * LAYOUT_COLUMN_COUNT),
            ),
        )
        mapped.setdefault(row_index, []).append((start_column, end_column, block))

    for row_index, row_blocks in mapped.items():
        ordered = sorted(row_blocks, key=lambda item: (item[0], item[2].x0))
        occupied_until = 0
        for block_index, (start_column, natural_end, block) in enumerate(ordered):
            start_column = max(start_column, occupied_until + 1)
            if start_column > LAYOUT_COLUMN_COUNT:
                continue
            next_start = (
                ordered[block_index + 1][0]
                if block_index + 1 < len(ordered)
                else LAYOUT_COLUMN_COUNT + 1
            )
            end_column = min(max(start_column, natural_end), max(start_column, next_start - 1))
            end_column = min(end_column, LAYOUT_COLUMN_COUNT)
            occupied_until = end_column

            cell = worksheet.cell(row=row_index, column=start_column, value=block.text)
            cell.number_format = "@"
            cell.quotePrefix = True
            cell.font = Font(
                name="Arial",
                size=max(8, min(20, block.font_size)),
                bold=block.bold,
                color="111827",
            )
            short_block = len(block.text) <= 24 and "\n" not in block.text
            cell.alignment = Alignment(
                vertical="top",
                horizontal="center" if block.text.strip().lower() == "release order" else "left",
                wrap_text=not short_block,
                shrink_to_fit=short_block,
            )
            if end_column > start_column:
                worksheet.merge_cells(
                    start_row=row_index,
                    start_column=start_column,
                    end_row=row_index,
                    end_column=end_column,
                )

            span_width = max(1, end_column - start_column + 1) * 3
            estimated_lines = 1 if short_block else max(1, math.ceil(len(block.text) / span_width))
            worksheet.row_dimensions[row_index].height = max(
                worksheet.row_dimensions[row_index].height or base_row_height,
                min(42, max(block.font_size * 1.35, estimated_lines * block.font_size * 1.2)),
            )


def _build_workbook(
    *,
    source_file_name: str,
    sheets: list[ExtractedSheet],
    page_count: int,
    table_count: int,
    text_page_count: int,
    ocr_page_count: int,
) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "转换说明"
    summary.sheet_view.showGridLines = False
    summary.append(["PDF 转 Excel", "转换结果"])
    summary.append(["源文件", Path(source_file_name).name])
    summary.append(["PDF 页数", str(page_count)])
    summary.append(["识别表格", str(table_count)])
    summary.append(["文字页", str(text_page_count)])
    summary.append(["OCR 页", str(ocr_page_count)])
    summary.append(["数据工作表", str(len(sheets))])
    summary.append(["转换策略", "有真实表格边界时提取为数据表；订单、表单等版式页面按原始坐标还原。"])
    summary.append(["说明", "源 PDF 未被修改；乱码字体页会自动尝试 OCR。OCR 结果建议与原 PDF 复核，所有单元格按文本保存以保留前导零。"])
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=12)
        cell.fill = PatternFill("solid", fgColor="0F766E")
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 88
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    existing_titles = {summary.title}
    for extracted in sheets:
        worksheet = workbook.create_sheet(_safe_sheet_title(extracted.title, existing_titles))
        if extracted.layout_blocks:
            _style_layout_sheet(worksheet, extracted)
        else:
            for row in extracted.rows:
                worksheet.append(row)
            _style_data_sheet(worksheet, extracted.rows)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _output_file_name(source_file_name: str) -> str:
    source_name = Path(source_file_name or "PDF文件.pdf").name
    stem = re.sub(r"[\\/:*?\"<>|]+", "_", Path(source_name).stem).strip(" .") or "PDF文件"
    return f"{stem}_转换结果.xlsx"


def convert_pdf_to_excel(pdf_bytes: bytes, source_file_name: str) -> PdfToExcelResult:
    try:
        document = pdfplumber.open(BytesIO(pdf_bytes))
    except Exception as exc:
        raise PdfToExcelConversionError("PDF 无法读取；请确认文件未加密、未损坏。") from exc

    sheets: list[ExtractedSheet] = []
    table_count = 0
    text_page_count = 0
    ocr_page_count = 0

    try:
        page_count = len(document.pages)
        if page_count == 0:
            raise PdfToExcelConversionError("PDF 中没有可转换的页面。")
        if page_count > MAX_PDF_PAGES:
            raise PdfToExcelConversionError(f"单次最多转换 {MAX_PDF_PAGES} 页 PDF。")

        for page_index, page in enumerate(document.pages):
            page_number = page_index + 1
            page_sheets, native_text = _extract_native_page(page, page_number)
            if page_sheets:
                sheets.extend(page_sheets)
                table_count += sum(sheet.source == "table" for sheet in page_sheets)
                continue

            page_width = float(getattr(page, "width", 0) or 0)
            page_height = float(getattr(page, "height", 0) or 0)
            native_text_unreliable = _native_text_is_unreliable(native_text)
            if native_text.strip() and not native_text_unreliable:
                native_layout = _extract_native_layout(page)
                if native_layout:
                    sheets.append(ExtractedSheet(
                        f"第{page_number}页_版式",
                        [],
                        "layout",
                        tuple(native_layout),
                        page_width,
                        page_height,
                    ))
                    text_page_count += 1
                    continue

            if native_text_unreliable or not native_text.strip():
                ocr_layout = _ocr_page_layout(
                    pdf_bytes,
                    page_index,
                    page_width,
                    page_height,
                )
                if ocr_layout:
                    sheets.append(ExtractedSheet(
                        f"第{page_number}页_OCR版式",
                        [],
                        "ocr_layout",
                        tuple(ocr_layout),
                        page_width,
                        page_height,
                    ))
                    ocr_page_count += 1
                    continue

            native_rows = _rows_from_text(native_text)
            if native_rows:
                sheets.append(ExtractedSheet(f"第{page_number}页_文字", native_rows, "text"))
                text_page_count += 1
                continue

            ocr_rows = _ocr_page(pdf_bytes, page_index)
            if ocr_rows:
                sheets.append(ExtractedSheet(f"第{page_number}页_OCR", ocr_rows, "ocr"))
                ocr_page_count += 1
    finally:
        document.close()

    if not sheets:
        raise PdfToExcelConversionError(
            "未识别到可写入 Excel 的文字或表格；如果是扫描件，请确认服务器已安装 OCR 引擎并尝试更清晰的 PDF。",
        )

    output_file_name = _output_file_name(source_file_name)
    content = _build_workbook(
        source_file_name=source_file_name,
        sheets=sheets,
        page_count=page_count,
        table_count=table_count,
        text_page_count=text_page_count,
        ocr_page_count=ocr_page_count,
    )
    return PdfToExcelResult(
        content=content,
        page_count=page_count,
        table_count=table_count,
        text_page_count=text_page_count,
        ocr_page_count=ocr_page_count,
        output_file_name=output_file_name,
    )
