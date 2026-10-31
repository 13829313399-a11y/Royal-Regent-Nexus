"""New sample/misc demands: bounded workbook preview and atomic draft creation.

The original workbook is re-read before confirmation. Corrections are bound to
real source rows; no inventory, supplier acceptance or historical import states
are inferred from a worksheet.
"""
import hashlib
import json
from collections import OrderedDict
from datetime import date, datetime
from io import BytesIO
from zipfile import ZipFile, BadZipFile

from fastapi import HTTPException
from pydantic import BaseModel, ConfigDict, Field, ValidationError
from sqlalchemy import select

from app.models.carton_procurement import CartonAuditEvent
from app.schemas.carton_special_requests import Command, Create, Line
from app.services import carton_special_requests as special, carton_procurement as core

HEADERS = ["分组编号", "订单类型", "货号", "纸品类型", "纸质", "规格", "尺寸单位", "需求数量", "数量单位", "需用日期", "备注"]
KEYS = ["group", "kind", "item_no", "packaging_type", "paper_quality", "specification", "dimension_unit", "quantity", "unit", "required_date", "note"]
MAX_BYTES = 5 * 1024 * 1024


class ImportRow(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)
    source_row: int = Field(ge=2, le=551)
    group: str = Field(min_length=1, max_length=64)
    kind: str = Field(max_length=32)
    item_no: str = Field(default="", max_length=128)
    packaging_type: str = Field(max_length=64)
    paper_quality: str = Field(max_length=128)
    specification: str = Field(max_length=255)
    dimension_unit: str = Field(max_length=16)
    quantity: str = Field(max_length=32)
    unit: str = Field(max_length=32)
    required_date: str = Field(max_length=10)
    note: str = Field(default="", max_length=2000)


class Confirm(Command):
    fingerprint: str = Field(pattern=r"^[a-f0-9]{64}$")
    rows: list[ImportRow] = Field(min_length=1, max_length=500)


def parse(content):
    """Runs in the disposable file worker, never against the business database."""
    from openpyxl import load_workbook
    if not content or len(content) > MAX_BYTES:
        raise HTTPException(422, "请选择不超过 5 MiB 的新建需求模板 .xlsx 文件")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 500 or sum(x.file_size for x in archive.infolist()) > 20 * 1024 * 1024:
                raise HTTPException(422, "文件展开后过大，请使用新建需求模板并缩小批次")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, ValueError, KeyError, OSError) as exc:
        raise HTTPException(422, "无法读取文件，请使用新建需求模板 .xlsx") from exc
    try:
        if "新建需求" not in book.sheetnames:
            raise HTTPException(422, "请下载新建需求模板填写；旧样办表不能直接作为本模板导入")
        sheet = book["新建需求"]
        if (sheet.max_row or 0) > 501 or (sheet.max_column or 0) > len(HEADERS):
            raise HTTPException(422, "每次最多 500 行，且不能添加模板以外的列")
        cells = list(sheet.iter_rows(max_col=len(HEADERS)))
        if not cells or [str(c.value or "").strip() for c in cells[0]] != HEADERS:
            raise HTTPException(422, "模板表头不匹配，请使用最新的新建需求模板")
        rows = []
        for row_no, cells_row in enumerate(cells[1:], 2):
            if not any(c.value is not None for c in cells_row):
                continue
            values = []
            for cell in cells_row:
                if cell.data_type in {"f", "e"}:
                    raise HTTPException(422, f"第 {row_no} 行含公式或错误值，请粘贴为实际数值")
                value = cell.value
                if isinstance(value, (date, datetime)): value = value.strftime("%Y-%m-%d")
                text = "" if value is None else str(value).strip()
                if len(text) > 2000: raise HTTPException(422, f"第 {row_no} 行内容过长")
                values.append(text)
            rows.append({"source_row": row_no, **dict(zip(KEYS, values))})
        if not rows: raise HTTPException(422, "文件没有需求明细，请填写后再导入")
        return {"fingerprint": hashlib.sha256(content).hexdigest(), "rows": rows}
    finally:
        book.close()


def authorize(db, user, factory):
    special.access(db, user, factory)
    if not special.can_submit(user, factory): raise HTTPException(403, "没有本厂样板／杂单提交权限")


def grouped(rows):
    groups = OrderedDict()
    for row in rows:
        kind = {"样板单": "SAMPLE", "样板": "SAMPLE", "杂单": "WORKSHOP_MISC", "SAMPLE": "SAMPLE", "WORKSHOP_MISC": "WORKSHOP_MISC"}.get(row.kind)
        if not kind: raise HTTPException(422, f"第 {row.source_row} 行：订单类型请选择样板单或杂单")
        dimension = {"cm": "cm", "厘米": "cm", "in": "in", "inch": "in", "英寸": "in"}.get(row.dimension_unit.lower())
        if not dimension: raise HTTPException(422, f"第 {row.source_row} 行：尺寸单位须为 cm 或 inch")
        try:
            paper = Line(item_no=row.item_no, packaging_type=row.packaging_type, paper_quality=row.paper_quality,
                specification=row.specification, dimension_unit=dimension, quantity=row.quantity, unit=row.unit)
            required = date.fromisoformat(row.required_date)
            core._unordered_specification(paper.specification, paper.dimension_unit)
        except (ValueError, ValidationError, HTTPException) as exc:
            raise HTTPException(422, f"第 {row.source_row} 行：请核对纸品、规格单位、正数数量及需用日期（YYYY-MM-DD）") from exc
        item = groups.setdefault(row.group, {"kind": kind, "required_date": required, "lines": [], "notes": [], "source_rows": []})
        if (item["kind"], item["required_date"]) != (kind, required):
            raise HTTPException(422, f"分组 {row.group} 的订单类型和需用日期必须一致，请修改或拆为不同分组")
        item["lines"].append(paper)
        item["source_rows"].append(row.source_row)
        if row.note and row.note not in item["notes"]: item["notes"].append(row.note)
        if len(item["lines"]) > 50: raise HTTPException(422, f"分组 {row.group} 超过 50 项纸品，请在文件中明确分组")
        if len("\n".join(item["notes"])) > 2000: raise HTTPException(422, f"分组 {row.group} 备注合计不能超过 2000 字")
    if len(groups) > 100: raise HTTPException(422, "每次最多新建 100 张订单")
    return groups


def confirm(db, user, payload, original, filename, *, refresh_actor):
    if original["fingerprint"] != payload.fingerprint:
        raise HTTPException(409, "文件已变化，请重新预览")
    source_ids = {row["source_row"] for row in original["rows"]}
    selected_ids = [row.source_row for row in payload.rows]
    if len(set(selected_ids)) != len(selected_ids) or set(selected_ids) - source_ids:
        raise HTTPException(422, "导入行重复或不属于此文件，请重新预览")
    user = special.locked_actor(db, user, payload.factory_id)
    user = refresh_actor()
    authorize(db, user, payload.factory_id)
    digest, saved = special.replay(db, user, payload, "DEMAND_IMPORT")
    if saved:
        for row in saved["orders"]: special.access(db, user, payload.factory_id, row["id"])
        return saved
    source_groups = {}
    for row in original["rows"]:
        source_groups.setdefault(row.get("group", "").strip(), set()).add(row["source_row"])
    selected = set(selected_ids)
    confirmed_groups = {row.source_row: row.group for row in payload.rows}
    for group, members in source_groups.items():
        if not group and selected != source_ids:
            raise HTTPException(422, "原表有未明确分组的纸品，请补齐全部订单归属后导入")
        if members & selected and (not members <= selected or (group and len({confirmed_groups[n] for n in members}) != 1)):
            raise HTTPException(422, "同一订单全部纸品须一起导入，不能只选择部分原表行")
    groups = grouped(payload.rows)
    try:
        # One immutable source-row claim prevents repeat files and changed group labels
        # from re-importing a row, including under another authorized operator.
        claims = {n: "CAE-DI-" + hashlib.sha256(f"{payload.factory_id}:{payload.fingerprint}:{n}".encode()).hexdigest() for n in selected_ids}
        if db.scalar(select(CartonAuditEvent.id).where(CartonAuditEvent.id.in_(claims.values())).limit(1)):
            raise HTTPException(409, "所选文件行已有导入记录，请核对订单台账，不能重复下单")
        orders = []
        for group, item in groups.items():
            key = "DI-" + hashlib.sha256(f"{payload.request_id}:{group}".encode()).hexdigest()
            body = Create(factory_id=payload.factory_id, request_id=key, kind=item["kind"],
                required_date=item["required_date"], purpose="\n".join(item["notes"]), lines=item["lines"])
            created = special.create(db, user, body, commit=False)
            orders.append({"group": group, "id": created["id"], "number": created["request_no"], "kind": created["kind"], "source_rows": item["source_rows"]})
            for n in item["source_rows"]:
                event = core._audit(db, user, payload.factory_id, "DEMAND_IMPORTED", "carton_special_request", created["id"],
                    {"batch_id": payload.request_id, "filename": filename, "sha256": payload.fingerprint,
                     "sheet": "新建需求", "source_row": n,
                     "original": next(row for row in original["rows"] if row["source_row"] == n),
                     "confirmed": next(row.model_dump() for row in payload.rows if row.source_row == n)})
                event.id = claims[n]
            db.flush()
        result = {"batch_id": payload.request_id, "count": len(orders), "orders": orders}
        event = core._audit(db, user, payload.factory_id, "DEMAND_IMPORT", "carton_unified_order_batch", payload.request_id,
            {"fingerprint": digest, "response": result, "filename": filename})
        event.id = special.replay_id(user, payload)
        db.commit()
        return result
    except Exception:
        db.rollback()
        raise
