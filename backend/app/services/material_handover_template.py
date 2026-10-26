"""Print options for the known supplier workbook; never edit source bytes.

Unfamiliar supplier layouts deliberately retain original-layout printing.
"""
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

HEADERS = ('物料分类', '物料编码', '物料名称', '规格', '颜色', '数量', '单位', '缸号', '卷号', '生产单号', '单价', '币种', '备注')
SHEET = '送货单填写'


def print_options(content):
    """Recognize both existing 10/30-row and extended standard templates.

    Only completely unused material rows may disappear. Default CNY alone is
    unused; zero, formulas, partial identities, notes and foreign currency stay.
    Extra sheets, drawings or content outside the contract retain their layout.
    """
    with ZipFile(BytesIO(content)) as archive:
        if any(name.startswith('xl/drawings/') for name in archive.namelist()):
            return {}
    book = None
    try:
        book = load_workbook(BytesIO(content), data_only=False, keep_links=False)
        if book.sheetnames != [SHEET, '填写说明'] or book['填写说明']['A2'].value != '供应商填写说明':
            return {}
        sheet = book[SHEET]
        if sheet.max_column != 13 or not 14 <= sheet.max_row <= 220:
            return {}
        if sheet['A2'].value != '统一送货单（布料、辅料、线）' or tuple(cell.value for cell in sheet[7]) != HEADERS:
            return {}
        footers = [row for row in range(9, sheet.max_row + 1) if sheet.cell(row, 1).value == '签收信息']
        if len(footers) != 1:
            return {}
        footer = footers[0]
        if footer + 5 != sheet.max_row:
            return {}
        labels = {(footer + 1, 1): '开单员：', (footer + 1, 6): '送货人：',
                  (footer + 1, 10): '收货人：', (footer + 3, 1): '签收日期：', (footer + 3, 6): '签收备注：'}
        if any(sheet.cell(row, column).value != label for (row, column), label in labels.items()):
            return {}
        hidden = []
        for row in range(8, footer):
            used = False
            for column in range(1, 14):
                cell = sheet.cell(row, column)
                value = cell.value
                if cell.data_type != 'f' and (value is None or isinstance(value, str) and not value.strip()):
                    continue
                if column == 12 and cell.data_type != 'f' and value == 'CNY':
                    continue
                used = True
            if not used:
                hidden.append(row - 1)  # UNO uses zero-based indices.
        return {'sheets': [SHEET], 'range': f'A1:M{sheet.max_row}', 'paper': 'A4',
                'orientation': 'landscape', 'print_mode': 'fit_width',
                'compact_handover': {'sheet': SHEET, 'hidden_rows': hidden, 'footer_row': footer - 1}}
    except (BadZipFile, InvalidFileException, KeyError, ValueError):
        return {}
    finally:
        if book is not None:
            book.close()
