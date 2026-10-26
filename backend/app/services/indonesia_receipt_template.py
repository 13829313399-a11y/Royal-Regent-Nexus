"""Read-only order-scoped receipt templates; reference quantities never post receipts."""
from collections import defaultdict
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.indonesia_material_calculations import ZERO, decimal_text, number, precise_decimals, require
from app.services.indonesia_materials import Ledger, authorize, now_text, ready

TITLE = '仓库实收回传模板（订单物料）'
REFERENCE_HEADERS = {'本次订单物料需求数量', '已实收并分配数量', '剩余未送货数量'}
HEADERS = ['厂区', '客户', '接收方', '合同号', '客户PO', '货号', '交付批次', '物料编码',
           '正式报关中文名称', '正式报关英文名称', '规格', '数量单位',
           '本次订单物料需求数量', '已实收并分配数量', '剩余未送货数量',
           '仓库', '收料单号', '行号', '实收日期', '正常实收数量', '备品数量',
           '箱号', '箱数', '长cm', '宽cm', '高cm', '每箱纸箱重量kg']
ENTRY_FIELDS = ('receipt_no', 'receipt_line', 'actual_date', 'normal_quantity', 'quantity', 'spare_quantity',
                'package_label', 'cartons', 'length_cm', 'width_cm', 'height_cm', 'carton_weight_kg')
SHEET = '实收回传'
HEADER_ROW = 5
FIRST_ROW = HEADER_ROW + 1


def literal(cell, value):
    """Store even formula-looking names/codes as plain text, with no precision loss."""
    cell.value = str(value) if value is not None else None
    if value is not None:
        cell.data_type = 's'


@precise_decimals
def selected_rows(db, user, request):
    authorize(user, request.factory_id, 'write')
    ready(db)
    keys = [(row.order_id, row.demand_id) for row in request.selections]
    require(len(set(keys)) == len(keys), '不能重复选择同一合同物料', 422)
    ledger = Ledger(db, request.factory_id)
    supplies = {row.id: row for row in ledger.rows('supply')}
    received = defaultdict(lambda: ZERO)
    selected_keys = set(keys)
    for allocation in ledger.rows('allocation'):
        data = allocation.data
        key = data['order_id'], data['demand_id']
        if key not in selected_keys:
            continue
        source = supplies.get(data['supply_id'])
        require(source is not None, '合同分配的回料来源已不可用，请先核对台账')
        if source.data['status'] == 'received':
            received[key] += number(data['normal_quantity']) + number(data['spare_quantity'])
    result = []
    for choice in request.selections:
        order = ledger.get(choice.order_id, 'order')
        require(order.data['customer'] == request.customer and order.data['recipient'] == request.recipient,
                '所选合同不属于本次客人或接收方，请重新选择', 422)
        demand = next((item for item in order.data['demands'] if item['id'] == choice.demand_id), None)
        require(demand is not None, '合同物料需求已变更，请刷新后重新选择', 409)
        material = ledger.get(demand['material_id'], 'material').data
        require(material['unit'].casefold().strip() == demand['unit'].casefold().strip(),
                '合同与基础资料的数量单位不一致，请核对后下载', 409)
        quantity = number(demand['quantity'])
        coverage = received[(choice.order_id, choice.demand_id)]
        result.append([request.factory_id, request.customer, request.recipient, order.data['contract_no'],
            order.data.get('po_no', ''), order.data['item_no'], order.data.get('batch', ''), material['code'],
            material['name'], material.get('english_name', ''), material.get('specification', ''), material['unit'],
            decimal_text(quantity), decimal_text(coverage), decimal_text(max(quantity - coverage, ZERO)),
            request.warehouse, *([None] * 11)])
    return result


def build_template(db, user, request):
    rows = selected_rows(db, user, request)
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.sheet_view.showGridLines = False
    end = get_column_letter(len(HEADERS))
    notes = [TITLE, f'生成时间：{now_text()} · 灰色为系统资料，浅黄色为仓库填写区。',
             '参考数量为下载时快照；剩余未送货＝订单需求－已实收且分配的正常量及备品，最低为 0；后续出货不重复扣减。',
             '填写收料单号、原单行号、日期及本次正常／备品数量；未填写实收的物料行可留空。装箱资料已知则填，采购资料由业务导入时核对。']
    for index, note in enumerate(notes, 1):
        sheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=len(HEADERS))
        cell = sheet.cell(index, 1)
        literal(cell, note)
        cell.font = Font(name='微软雅黑', size=15 if index == 1 else 11, bold=index == 1, color='123D38' if index == 1 else '526373')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        sheet.row_dimensions[index].height = 30 if index == 1 else 24
    thin = Side(style='thin', color='D7E1E5')
    for column, header in enumerate(HEADERS, 1):
        cell = sheet.cell(HEADER_ROW, column)
        literal(cell, header)
        cell.font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF' if column <= 15 else '5B4300')
        cell.fill = PatternFill('solid', fgColor='286F66' if column <= 15 else 'F3D98D')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    widths = [13, 18, 11, 23, 23, 16, 16, 19, 26, 32, 24, 12, 19, 19, 19, 17, 23, 12, 17, 18, 17, 19, 13, 13, 13, 13, 20]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    # Keep matching context in the file without making warehouse staff handle it.
    for column in ('D', 'E', 'G'):
        sheet.column_dimensions[column].hidden = True
    sheet.row_dimensions[HEADER_ROW].height = 38
    for index, values in enumerate(rows, FIRST_ROW):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(index, column)
            literal(cell, value)
            cell.font = Font(name='微软雅黑', size=11, color='18333C')
            cell.fill = PatternFill('solid', fgColor='F0F3F6' if column <= 15 else 'FFF4CC')
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical='center', wrap_text=True, horizontal='right' if 13 <= column <= 15 else 'left')
            cell.number_format = '@'
            cell.protection = Protection(locked=column <= 15)
        sheet.row_dimensions[index].height = 38
    last_row = FIRST_ROW + len(rows) - 1
    sheet.freeze_panes = f'A{FIRST_ROW}'
    sheet.auto_filter.ref = f'A{HEADER_ROW}:{end}{last_row}'
    sheet.print_title_rows = f'1:{HEADER_ROW}'
    sheet.print_area = f'A1:{end}{last_row}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    # No password: prevents accidental changes while allowing intentional local reformatting.
    sheet.protection.sheet = True
    sheet.protection.autoFilter = False
    sheet.protection.selectLockedCells = False
    sheet.protection.selectUnlockedCells = False
    warehouses = DataValidation(type='list', formula1='"来料仓,胶件仓,半成品仓"', allow_blank=True)
    warehouses.errorTitle = '仓库名称'; warehouses.error = '可选择已有仓库或填写原单仓库名称。'
    warehouses.showErrorMessage = False
    sheet.add_data_validation(warehouses)
    warehouses.add(f'P{FIRST_ROW}:P{last_row}')
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()
