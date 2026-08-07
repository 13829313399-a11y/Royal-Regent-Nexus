from __future__ import annotations

import hashlib
import hmac
import json
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any
from uuid import uuid4
from zipfile import BadZipFile

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.injection_scheduling import (
    InjectionSchedulingMachine,
    InjectionSchedulingMold,
)
from app.models.injection_scheduling_execution import (
    InjectionSchedulingAuditEvent,
    InjectionSchedulingOrder,
    InjectionSchedulingPlan,
    InjectionSchedulingPlanOrderState,
    InjectionSchedulingProgressAdjustment,
    InjectionSchedulingShiftReport,
    InjectionSchedulingTask,
)
from app.models.injection_scheduling_export import InjectionSchedulingExportAudit
from app.schemas.injection_scheduling_export import InjectionSchedulingExportRequest
from app.services.auth import AuthContext
from app.services.injection_scheduling_execution import _actor_name, _now
from app.services.injection_scheduling_profile_registry import (
    profile_revision_for_factory,
)
from app.services.injection_scheduling_profiles import (
    SYSTEM_STANDARD_EXPORT_PROFILE,
    ImportProfile,
    profile_header_fingerprint,
)
from app.services.injection_scheduling_projection import CALCULATION_VERSION

META_SHEET_NAME = "_SYSTEM_META"
META_MARKER = "ROYAL_REGENT_INJECTION_SCHEDULING_META_V1"
META_SCHEMA = "injection-scheduling-export-meta-v1"
META_STORAGE_FORMAT = "SEGMENTED_ROWS_V1"
META_ROW_START = 8
EXCEL_CELL_TEXT_LIMIT = 32_767


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), sort_keys=True)


def _hash(value: bytes | str) -> str:
    content = value.encode("utf-8") if isinstance(value, str) else value
    return hashlib.sha256(content).hexdigest()


def _payload_hash(value: Any) -> str:
    return _hash(_json(value))


def _signing_material() -> tuple[bytes, str]:
    key = settings.injection_scheduling_export_signing_key.encode("utf-8")
    if len(key) < 32:
        raise HTTPException(
            status_code=503,
            detail={
                "code": "EXPORT_SIGNING_KEY_NOT_CONFIGURED",
                "message": "注塑排产导出签名密钥未配置或少于 32 字节",
            },
        )
    key_id = settings.injection_scheduling_export_signing_key_id.strip()
    if not key_id:
        raise HTTPException(status_code=503, detail="导出 signing key id 未配置")
    return key, key_id


def _verification_key(key_id: str) -> bytes:
    current_key, current_key_id = _signing_material()
    if key_id == current_key_id:
        return current_key
    try:
        configured = json.loads(
            settings.injection_scheduling_export_verification_keys_json or "{}"
        )
    except json.JSONDecodeError:
        configured = {}
    historical = configured.get(key_id) if isinstance(configured, dict) else None
    if not isinstance(historical, str) or len(historical.encode("utf-8")) < 32:
        raise HTTPException(
            status_code=409,
            detail={
                "code": "SYSTEM_META_SIGNING_KEY_REVOKED",
                "message": "_SYSTEM_META 使用的签名 key 已撤销或不可用",
            },
        )
    return historical.encode("utf-8")


def _signature(manifest_json: str, key: bytes) -> str:
    return hmac.new(key, manifest_json.encode("utf-8"), hashlib.sha256).hexdigest()


def _safe_excel_value(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _excel_cell_value(converter: str, value: Any) -> Any:
    if value in (None, ""):
        return ""
    if converter == "datetime" and isinstance(value, str):
        try:
            return datetime.fromisoformat(value).replace(tzinfo=None)
        except ValueError:
            return _safe_excel_value(value)
    if converter == "date" and isinstance(value, str):
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return _safe_excel_value(value)
    return _safe_excel_value(value)


def _lineage(order: InjectionSchedulingOrder) -> dict[str, Any]:
    try:
        value = json.loads(order.lineage_json or "{}")
        return value if isinstance(value, dict) else {}
    except json.JSONDecodeError:
        return {}


def _find_profile(
    db: Session,
    *,
    factory_id: str,
    profile_id: str,
    profile_revision: int,
) -> ImportProfile:
    if (
        profile_id == SYSTEM_STANDARD_EXPORT_PROFILE.profile_id
        and profile_revision == SYSTEM_STANDARD_EXPORT_PROFILE.revision
        and factory_id in SYSTEM_STANDARD_EXPORT_PROFILE.factories
    ):
        return SYSTEM_STANDARD_EXPORT_PROFILE
    profile = profile_revision_for_factory(
        db,
        factory_id=factory_id,
        profile_id=profile_id,
        profile_revision=profile_revision,
        allowed_statuses=("ACTIVE", "RETIRED"),
    )
    if profile is None:
        raise HTTPException(
            status_code=409,
            detail={
                "code": "EXPORT_PROFILE_UNAVAILABLE",
                "message": "指定的 ACTIVE Profile revision 不属于当前厂区",
                "profile_id": profile_id,
                "profile_revision": profile_revision,
            },
        )
    return profile


def _export_profile(
    db: Session,
    *,
    plan: InjectionSchedulingPlan,
    payload: InjectionSchedulingExportRequest,
) -> ImportProfile:
    if payload.export_mode == "SOURCE_COMPATIBLE":
        if (
            plan.export_binding_source != "IMPORT_PROFILE"
            or not plan.export_profile_id
            or plan.export_profile_revision is None
        ):
            raise HTTPException(
                status_code=409,
                detail={
                    "code": "PLAN_EXPORT_PROFILE_NOT_BOUND",
                    "message": "该计划没有显式来源 Profile binding，不能兼容原表导出",
                },
            )
        if (
            payload.profile_id != plan.export_profile_id
            or payload.profile_revision != plan.export_profile_revision
        ):
            raise HTTPException(
                status_code=409,
                detail={
                    "code": "PLAN_EXPORT_PROFILE_MISMATCH",
                    "message": "请求 Profile 与计划锁定的导出 Profile 不一致",
                    "bound_profile_id": plan.export_profile_id,
                    "bound_profile_revision": plan.export_profile_revision,
                },
            )
        return _find_profile(
            db,
            factory_id=plan.factory_id,
            profile_id=plan.export_profile_id,
            profile_revision=plan.export_profile_revision,
        )
    return _find_profile(
        db,
        factory_id=plan.factory_id,
        profile_id="isprofile-system-standard-v1",
        profile_revision=1,
    )


def _canonical_rows(
    db: Session,
    *,
    plan: InjectionSchedulingPlan,
) -> list[dict[str, Any]]:
    tasks = list(
        db.scalars(
            select(InjectionSchedulingTask)
            .where(
                InjectionSchedulingTask.factory_id == plan.factory_id,
                InjectionSchedulingTask.plan_id == plan.id,
            )
            .order_by(
                InjectionSchedulingTask.machine_id,
                InjectionSchedulingTask.sequence_no,
                InjectionSchedulingTask.id,
            )
        ).all()
    )
    order_ids = {item.order_id for item in tasks}
    machine_ids = {item.machine_id for item in tasks}
    mold_ids = {item.mold_id for item in tasks if item.mold_id}
    orders = {
        item.id: item
        for item in db.scalars(
            select(InjectionSchedulingOrder).where(
                InjectionSchedulingOrder.factory_id == plan.factory_id,
                InjectionSchedulingOrder.id.in_(order_ids),
            )
        ).all()
    }
    machines = {
        item.id: item
        for item in db.scalars(
            select(InjectionSchedulingMachine).where(
                InjectionSchedulingMachine.factory_id == plan.factory_id,
                InjectionSchedulingMachine.id.in_(machine_ids),
            )
        ).all()
    }
    molds = (
        {
            item.id: item
            for item in db.scalars(
                select(InjectionSchedulingMold).where(
                    InjectionSchedulingMold.factory_id == plan.factory_id,
                    InjectionSchedulingMold.id.in_(mold_ids),
                )
            ).all()
        }
        if mold_ids
        else {}
    )
    states = {
        item.order_id: item
        for item in db.scalars(
            select(InjectionSchedulingPlanOrderState).where(
                InjectionSchedulingPlanOrderState.factory_id == plan.factory_id,
                InjectionSchedulingPlanOrderState.plan_id == plan.id,
            )
        ).all()
    }
    reports_by_task: dict[str, list[InjectionSchedulingShiftReport]] = {}
    if order_ids:
        for report in db.scalars(
            select(InjectionSchedulingShiftReport)
            .where(
                InjectionSchedulingShiftReport.factory_id == plan.factory_id,
                InjectionSchedulingShiftReport.task_id.in_({item.id for item in tasks}),
            )
            .order_by(
                InjectionSchedulingShiftReport.created_at,
                InjectionSchedulingShiftReport.id,
            )
        ).all():
            reports_by_task.setdefault(report.task_id, []).append(report)
    adjustments_by_order: dict[str, list[InjectionSchedulingProgressAdjustment]] = {}
    if order_ids:
        for adjustment in db.scalars(
            select(InjectionSchedulingProgressAdjustment)
            .where(
                InjectionSchedulingProgressAdjustment.factory_id == plan.factory_id,
                InjectionSchedulingProgressAdjustment.plan_id == plan.id,
                InjectionSchedulingProgressAdjustment.order_id.in_(order_ids),
            )
            .order_by(
                InjectionSchedulingProgressAdjustment.created_at,
                InjectionSchedulingProgressAdjustment.id,
            )
        ).all():
            adjustments_by_order.setdefault(adjustment.order_id, []).append(adjustment)
    rows: list[dict[str, Any]] = []
    for task in tasks:
        order = orders.get(task.order_id)
        machine = machines.get(task.machine_id)
        mold = molds.get(task.mold_id or "")
        if order is None or machine is None:
            raise HTTPException(status_code=409, detail="计划任务引用的订单或机台不存在")
        state = states.get(order.id)
        completed = float(state.completed_quantity if state is not None else order.completed_quantity)
        order_quantity = float(state.order_quantity if state is not None else order.order_quantity)
        lineage = _lineage(order)
        reports = reports_by_task.get(task.id, [])
        adjustments = adjustments_by_order.get(order.id, [])
        report_ledger_digest = _payload_hash(
            {
                "reports": [
                    {
                        "id": item.id,
                        "request_id": item.request_id,
                        "normalized_increment_quantity": str(
                            item.normalized_increment_quantity
                        ),
                        "created_at": item.created_at,
                    }
                    for item in reports
                ],
                "adjustments": [
                    {
                        "id": item.id,
                        "request_id": item.request_id,
                        "signed_quantity": str(item.signed_quantity),
                        "created_at": item.created_at,
                    }
                    for item in adjustments
                ],
            }
        )
        planned_quantity = float(task.allocated_quantity)
        if planned_quantity <= 0:
            planned_quantity = max(order_quantity - completed, 0)
        rows.append(
            {
                "machine_code": machine.machine_code,
                "automation_mode": lineage.get("automation", ""),
                "legacy_marker": "锁定" if task.locked else lineage.get("marker", ""),
                "legacy_machine_class_text": mold.mold_class_raw if mold else "",
                "mold_no": mold.mold_no if mold else "",
                "product_name": order.product_name,
                "order_no": order.order_no,
                "item_no": order.item_no,
                "warehouse_text": order.warehouse_text,
                "set_quantity": lineage.get("set_quantity", ""),
                "order_quantity": order_quantity,
                "completed_quantity": completed,
                "source_outstanding_quantity": max(order_quantity - completed, 0),
                "daily_target_quantity": lineage.get("daily_target_quantity", ""),
                "shift_target_quantity": float(task.shift_target_quantity),
                "material_name": mold.material_name if mold else "",
                "sprue_ratio": lineage.get("sprue_ratio", ""),
                "color_name": mold.color_profile if mold else "",
                "color_powder_code": lineage.get("powder", ""),
                "whole_shot_net_weight_g": float(mold.whole_shot_net_weight_g)
                if mold and mold.whole_shot_net_weight_g is not None
                else "",
                "whole_shot_gross_weight_g": float(mold.whole_shot_gross_weight_g)
                if mold and mold.whole_shot_gross_weight_g is not None
                else "",
                "material_weight_kg": lineage.get("material_kg", ""),
                "unit_price": lineage.get("unit_price", ""),
                "outsource_price": lineage.get("outsource_price", ""),
                "price_ratio": lineage.get("ratio", ""),
                "order_date": lineage.get("order_date", ""),
                "delivery_start_date": order.delivery_start_date,
                "delivery_due_date": order.delivery_due_date,
                "legacy_mold_change": lineage.get("mold_change_ref", ""),
                "legacy_color_change": lineage.get("color_change_ref", ""),
                "legacy_setup": lineage.get("setup_time", ""),
                "legacy_downtime": task.planned_downtime_minutes,
                "planned_start": task.planned_start,
                "planned_finish": task.planned_finish,
                "source_plan_month": task.planned_finish[:7],
                "source_warehouse_date": lineage.get("warehouse_date", ""),
                "source_delivery_slack": task.delivery_slack_days,
                "requires_spray_paint": lineage.get("spray", ""),
                "source_production_days": round(task.production_minutes / 1440, 4),
                "ship_date": lineage.get("ship_date", ""),
                "remark": order.remark,
                "required_arm_type": mold.required_arm_type if mold else "",
                "required_fixture_type": mold.required_fixture_type if mold else "",
                "source_machine_finish": task.estimated_finish,
                "stable_order_key": task.stable_order_key,
                "stable_row_key": task.stable_row_key,
                "split_key": task.stable_row_key,
                "planned_quantity": planned_quantity,
                "_machine_area": machine.area,
                "_machine_injection_capacity_g": float(machine.injection_capacity_g)
                if machine.injection_capacity_g is not None
                else "",
                "_machine_type": machine.machine_type,
                "_system_meta": {
                    "row_type": "TASK",
                    "machine_id": task.machine_id,
                    "order_id": task.order_id,
                    "task_id": task.id,
                    "stable_order_key": task.stable_order_key,
                    "stable_row_key": task.stable_row_key,
                    "split_key": task.stable_row_key,
                    "quantity_scope": state.quantity_scope
                    if state is not None
                    else "ORDER_CUMULATIVE",
                    "planned_quantity": planned_quantity,
                    "order_revision": order.revision,
                    "task_revision": task.revision,
                    "task_report_counter_at_export": float(task.reported_quantity),
                    "takeover_source_completed_quantity": float(
                        state.takeover_source_completed_quantity
                        if state is not None
                        else task.takeover_source_completed_quantity
                    ),
                    "completed_at_export": completed,
                    "report_increment_total_at_export": float(
                        state.report_increment_total if state is not None else 0
                    ),
                    "progress_adjustment_total_at_export": float(
                        state.progress_adjustment_total if state is not None else 0
                    ),
                    "report_ledger_digest": report_ledger_digest,
                },
            }
        )
    return rows


def _profile_role_name(profile: ImportProfile, role: str) -> str:
    return next(
        (item.names[0] for item in profile.sheet_roles if item.role == role),
        "",
    )


def _dynamic_shift_values(
    profile: ImportProfile,
    rows: list[dict[str, Any]],
) -> tuple[list[tuple[str, str, str]], dict[str, dict[str, float]]]:
    if (
        profile.profile_family == "system_standard"
        or not profile.dynamic_shift_start
        or not profile.dynamic_shift_end
    ):
        return [], {}
    by_row: dict[str, dict[str, float]] = {}
    ordered_keys: dict[str, datetime] = {}
    for row in rows:
        try:
            cursor = datetime.fromisoformat(str(row.get("planned_start", ""))).replace(
                tzinfo=None
            )
            finish = datetime.fromisoformat(
                str(row.get("planned_finish", ""))
            ).replace(tzinfo=None)
        except ValueError:
            continue
        remaining = float(row.get("planned_quantity") or 0)
        target = float(row.get("shift_target_quantity") or 0) or remaining
        allocations: dict[str, float] = {}
        for _ in range(370):
            if cursor >= finish or remaining <= 0:
                break
            shift_code = "DAY" if 8 <= cursor.hour < 20 else "NIGHT"
            key = f"{cursor.date().isoformat()}:{shift_code}"
            allocations[key] = allocations.get(key, 0) + min(target, remaining)
            ordered_keys.setdefault(key, cursor)
            remaining -= min(target, remaining)
            cursor += timedelta(hours=12)
        by_row[str(row.get("stable_row_key", ""))] = allocations
    start_index = column_index_from_string(profile.dynamic_shift_start)
    end_index = column_index_from_string(profile.dynamic_shift_end)
    ordered = sorted(ordered_keys, key=lambda item: (ordered_keys[item], item))
    columns: list[tuple[str, str, str]] = []
    for offset, key in enumerate(ordered):
        column_index = start_index + offset
        if column_index > end_index:
            break
        day, shift_code = key.split(":", 1)
        parsed_day = date.fromisoformat(day)
        label = f"{'白班' if shift_code == 'DAY' else '夜班'} {parsed_day.month}/{parsed_day.day}"
        columns.append((get_column_letter(column_index), label, key))
    return columns, by_row


def _render_profile_master_sheets(
    workbook: Workbook,
    *,
    profile: ImportProfile,
    rows: list[dict[str, Any]],
) -> None:
    mold_sheet_name = _profile_role_name(profile, "MOLD_MASTER")
    if mold_sheet_name:
        sheet = workbook.create_sheet(mold_sheet_name)
        headers = {
            "A": "模号",
            "B": "名称",
            "G": "安数",
            "H": "单双臂",
            "I": "吸盘",
            "L": "用料",
            "M": "工程重量（G）",
        }
        for column, value in headers.items():
            sheet[f"{column}1"] = value
        seen: set[str] = set()
        row_number = 2
        for row in rows:
            mold_no = str(row.get("mold_no", ""))
            if not mold_no or mold_no in seen:
                continue
            seen.add(mold_no)
            values = {
                "A": mold_no,
                "B": row.get("product_name", ""),
                "G": row.get("legacy_machine_class_text", ""),
                "H": row.get("required_arm_type", ""),
                "I": row.get("required_fixture_type", ""),
                "L": row.get("material_name", ""),
                "M": row.get("whole_shot_net_weight_g", ""),
            }
            for column, value in values.items():
                sheet[f"{column}{row_number}"] = _safe_excel_value(value)
            row_number += 1

    machine_sheet_name = _profile_role_name(profile, "MACHINE_MASTER")
    if not machine_sheet_name:
        return
    sheet = workbook.create_sheet(machine_sheet_name)
    huakang = profile.profile_family == "huakang_b_daily_plan"
    header_row = 3 if huakang else 2
    headers = (
        {
            "A": "设备编号",
            "B": "设备名称",
            "D": "安数",
            "F": "射胶量",
            "I": "设备类型",
            "J": "机械手",
        }
        if huakang
        else {
            "A": "摆放区域",
            "B": "机位",
            "E": "安数",
            "H": "射胶量",
            "L": "设备类型",
            "M": "机械手",
        }
    )
    for column, value in headers.items():
        sheet[f"{column}{header_row}"] = value
    seen = set()
    row_number = header_row + 1
    for row in rows:
        machine_code = str(row.get("machine_code", ""))
        if not machine_code or machine_code in seen:
            continue
        seen.add(machine_code)
        values = (
            {
                "A": machine_code,
                "B": machine_code,
                "D": row.get("legacy_machine_class_text", ""),
                "F": row.get("_machine_injection_capacity_g", ""),
                "I": row.get("_machine_type", ""),
                "J": row.get("automation_mode", ""),
            }
            if huakang
            else {
                "A": row.get("_machine_area", ""),
                "B": machine_code,
                "E": row.get("legacy_machine_class_text", ""),
                "H": row.get("_machine_injection_capacity_g", ""),
                "L": row.get("_machine_type", ""),
                "M": row.get("automation_mode", ""),
            }
        )
        for column, value in values.items():
            sheet[f"{column}{row_number}"] = _safe_excel_value(value)
        row_number += 1


def _render_workbook(
    *,
    profile: ImportProfile,
    rows: list[dict[str, Any]],
    manifest_json: str,
    signature: str,
    key_id: str,
) -> bytes:
    parsed_manifest = json.loads(manifest_json)
    row_manifests = parsed_manifest.pop("rows")
    global_manifest_json = _json(parsed_manifest)
    row_manifest_json = [_json(item) for item in row_manifests]
    if len(global_manifest_json) > EXCEL_CELL_TEXT_LIMIT or any(
        len(item) > EXCEL_CELL_TEXT_LIMIT for item in row_manifest_json
    ):
        raise HTTPException(
            status_code=409,
            detail={
                "code": "SYSTEM_META_CELL_LIMIT_EXCEEDED",
                "message": "签名导出元数据超过 Excel 单元格安全上限",
            },
        )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = profile.current_plan_role.names[0]
    header_row = profile.current_plan_role.header_row or 1
    if header_row > 1:
        sheet.cell(1, 1, profile.name)
        sheet.cell(1, 1).font = Font(bold=True, size=14)
    for field in profile.fields:
        cell = sheet[f"{field.column}{header_row}"]
        cell.value = field.headers[0]
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(horizontal="center")
    for row_number, source in enumerate(rows, start=header_row + 1):
        for field in profile.fields:
            value = _excel_cell_value(
                field.converter,
                source.get(field.canonical_field, ""),
            )
            cell = sheet[f"{field.column}{row_number}"]
            cell.value = value
            if field.converter == "identifier":
                cell.number_format = "@"
            elif field.converter == "datetime":
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif field.converter == "date":
                cell.number_format = "yyyy-mm-dd"
    dynamic_columns, dynamic_values = _dynamic_shift_values(profile, rows)
    for column, label, _key in dynamic_columns:
        cell = sheet[f"{column}{header_row}"]
        cell.value = label
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(horizontal="center")
    for row_number, source in enumerate(rows, start=header_row + 1):
        allocations = dynamic_values.get(str(source.get("stable_row_key", "")), {})
        for column, _label, key in dynamic_columns:
            sheet[f"{column}{row_number}"] = allocations.get(key, "")
    sheet.freeze_panes = f"A{header_row + 1}"
    final_column = dynamic_columns[-1][0] if dynamic_columns else profile.fields[-1].column
    sheet.auto_filter.ref = f"A{header_row}:{final_column}{header_row + len(rows)}"
    sheet.sheet_view.showGridLines = False
    for field in profile.fields:
        sheet.column_dimensions[field.column].width = max(12, min(28, len(field.headers[0]) * 2 + 4))
    for column, _label, _key in dynamic_columns:
        sheet.column_dimensions[column].width = 13
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:{final_column}{header_row + len(rows)}"

    _render_profile_master_sheets(workbook, profile=profile, rows=rows)

    meta = workbook.create_sheet(META_SHEET_NAME)
    meta["A1"] = META_MARKER
    meta["A2"] = global_manifest_json
    meta["A3"] = signature
    meta["A4"] = key_id
    meta["A5"] = _hash(manifest_json)
    meta["A6"] = META_STORAGE_FORMAT
    meta["A7"] = "ROW_MANIFEST_JSON"
    for offset, row_json in enumerate(row_manifest_json):
        meta.cell(META_ROW_START + offset, 1, row_json)
    meta.sheet_state = "veryHidden"
    fixed_time = datetime(2000, 1, 1)  # noqa: DTZ001
    workbook.properties.created = fixed_time
    workbook.properties.modified = fixed_time
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@dataclass(frozen=True, slots=True)
class ExportedWorkbook:
    audit_id: str
    file_name: str
    content: bytes
    file_sha256: str
    export_mode: str
    profile_id: str | None
    profile_revision: int
    renderer_code: str
    signing_key_id: str
    idempotent_replay: bool


def export_plan_workbook(
    db: Session,
    *,
    plan_id: str,
    payload: InjectionSchedulingExportRequest,
    user: AuthContext,
) -> ExportedWorkbook:
    request_payload_hash = _payload_hash(payload.model_dump(mode="json") | {"plan_id": plan_id})
    replay = db.scalar(
        select(InjectionSchedulingExportAudit).where(
            InjectionSchedulingExportAudit.factory_id == payload.factory_id,
            InjectionSchedulingExportAudit.request_id == payload.request_id,
        )
    )
    if replay is not None:
        if replay.plan_id != plan_id or replay.request_payload_hash != request_payload_hash:
            raise HTTPException(status_code=409, detail="相同导出 request_id 已用于其他请求")
        return ExportedWorkbook(
            audit_id=replay.id,
            file_name=replay.file_name,
            content=bytes(replay.payload_blob),
            file_sha256=replay.file_sha256,
            export_mode=replay.export_mode,
            profile_id=replay.profile_id,
            profile_revision=replay.profile_revision,
            renderer_code=replay.renderer_code,
            signing_key_id=replay.signing_key_id,
            idempotent_replay=True,
        )
    plan = db.scalar(
        select(InjectionSchedulingPlan).where(
            InjectionSchedulingPlan.id == plan_id,
            InjectionSchedulingPlan.factory_id == payload.factory_id,
        )
    )
    if plan is None:
        raise HTTPException(status_code=404, detail="排产计划不存在")
    if plan.revision != payload.expected_plan_revision:
        raise HTTPException(
            status_code=409,
            detail={
                "code": "PLAN_REVISION_STALE",
                "expected_revision": payload.expected_plan_revision,
                "current_revision": plan.revision,
            },
        )
    profile = _export_profile(db, plan=plan, payload=payload)
    rows = _canonical_rows(db, plan=plan)
    if not rows:
        raise HTTPException(status_code=409, detail="计划没有可导出的任务")
    key, key_id = _signing_material()
    audit_id = f"isexport-{uuid4().hex}"
    generated_at = _now()
    header_row = profile.current_plan_role.header_row or 1
    row_manifests = [
        {
            **row["_system_meta"],
            "source_row": header_row + index,
        }
        for index, row in enumerate(rows, start=1)
    ]
    signed_row_manifest_digest = _payload_hash(row_manifests)
    reference_report_event_sequence = int(
        db.scalar(
            select(func.max(InjectionSchedulingAuditEvent.sequence)).where(
                InjectionSchedulingAuditEvent.factory_id == plan.factory_id,
                InjectionSchedulingAuditEvent.event_type == "shift_report_recorded",
            )
        )
        or 0
    )
    mapping_fingerprint = profile_header_fingerprint(profile)
    manifest = {
        "schema": META_SCHEMA,
        "schema_version": META_SCHEMA,
        "audit_id": audit_id,
        "export_id": audit_id,
        "factory_id": plan.factory_id,
        "plan_id": plan.id,
        "plan_revision": plan.revision,
        "plan_status": plan.status,
        "rule_revision": plan.rule_revision,
        "export_mode": payload.export_mode,
        "profile_id": profile.profile_id,
        "profile_revision": profile.revision,
        "profile_family": profile.profile_family,
        "renderer_code": profile.renderer_code,
        "plan_export_profile_id": plan.export_profile_id,
        "plan_export_profile_revision": plan.export_profile_revision,
        "plan_export_profile_family": plan.export_profile_family,
        "plan_export_renderer_code": plan.export_renderer_code,
        "plan_export_binding_source": plan.export_binding_source,
        "mapping_fingerprint": mapping_fingerprint,
        "calculation_version": plan.calculation_version or CALCULATION_VERSION,
        "generated_at": generated_at,
        "exported_at": generated_at,
        "exported_by": user.id,
        "reference_report_event_sequence": reference_report_event_sequence,
        "task_count": len(rows),
        "signed_row_manifest_digest": signed_row_manifest_digest,
        "rows": row_manifests,
    }
    manifest_json = _json(manifest)
    signature = _signature(manifest_json, key)
    content = _render_workbook(
        profile=profile,
        rows=rows,
        manifest_json=manifest_json,
        signature=signature,
        key_id=key_id,
    )
    file_sha256 = _hash(content)
    mode_suffix = "source-compatible" if payload.export_mode == "SOURCE_COMPATIBLE" else "system-standard"
    file_name = f"injection-plan-{plan.factory_id}-{plan.business_date}-r{plan.revision}-{mode_suffix}.xlsx"
    audit = InjectionSchedulingExportAudit(
        id=audit_id,
        factory_id=plan.factory_id,
        plan_id=plan.id,
        plan_revision=plan.revision,
        export_mode=payload.export_mode,
        profile_id=profile.profile_id,
        profile_revision=profile.revision,
        profile_family=profile.profile_family,
        renderer_code=profile.renderer_code,
        calculation_version=manifest["calculation_version"],
        rule_revision=plan.rule_revision,
        mapping_fingerprint=mapping_fingerprint,
        reference_report_event_sequence=reference_report_event_sequence,
        signed_row_manifest_digest=signed_row_manifest_digest,
        export_options_json=_json({"export_mode": payload.export_mode}),
        file_name=file_name,
        file_sha256=file_sha256,
        size_bytes=len(content),
        manifest_json=manifest_json,
        manifest_sha256=_hash(manifest_json),
        metadata_signature=signature,
        signing_key_id=key_id,
        payload_blob=content,
        request_id=payload.request_id,
        request_payload_hash=request_payload_hash,
        created_by=user.id,
        created_by_name=_actor_name(user),
        created_at=generated_at,
    )
    db.add(audit)
    db.commit()
    return ExportedWorkbook(
        audit_id=audit.id,
        file_name=file_name,
        content=content,
        file_sha256=file_sha256,
        export_mode=payload.export_mode,
        profile_id=profile.profile_id,
        profile_revision=profile.revision,
        renderer_code=profile.renderer_code,
        signing_key_id=key_id,
        idempotent_replay=False,
    )


def verify_signed_system_meta(
    db: Session,
    content: bytes,
    *,
    factory_id: str,
) -> dict[str, Any] | None:
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except (BadZipFile, InvalidFileException, OSError, ValueError):
        return None
    if META_SHEET_NAME not in workbook.sheetnames:
        workbook.close()
        return None
    sheet = workbook[META_SHEET_NAME]
    marker = str(sheet["A1"].value or "")
    stored_manifest_json = str(sheet["A2"].value or "")
    supplied_signature = str(sheet["A3"].value or "")
    supplied_key_id = str(sheet["A4"].value or "")
    supplied_manifest_hash = str(sheet["A5"].value or "")
    storage_format = str(sheet["A6"].value or "")
    try:
        stored_manifest = json.loads(stored_manifest_json)
        if not isinstance(stored_manifest, dict):
            raise TypeError
        if storage_format == META_STORAGE_FORMAT:
            task_count = stored_manifest.get("task_count")
            if (
                not isinstance(task_count, int)
                or isinstance(task_count, bool)
                or task_count < 0
                or task_count > 100_000
            ):
                raise ValueError
            rows = []
            for offset in range(task_count):
                row_json = str(
                    sheet.cell(META_ROW_START + offset, 1).value or ""
                )
                row = json.loads(row_json)
                if not isinstance(row, dict):
                    raise TypeError
                rows.append(row)
            manifest = {**stored_manifest, "rows": rows}
            manifest_json = _json(manifest)
        else:
            manifest = stored_manifest
            manifest_json = _json(manifest)
    except (json.JSONDecodeError, TypeError, ValueError):
        workbook.close()
        raise HTTPException(
            status_code=409,
            detail={"code": "SYSTEM_META_INVALID", "message": "_SYSTEM_META 结构或摘要无效"},
        ) from None
    workbook.close()
    if marker != META_MARKER or _hash(manifest_json) != supplied_manifest_hash:
        raise HTTPException(
            status_code=409,
            detail={"code": "SYSTEM_META_INVALID", "message": "_SYSTEM_META 结构或摘要无效"},
        )
    key = _verification_key(supplied_key_id)
    if not hmac.compare_digest(
        supplied_signature,
        _signature(manifest_json, key),
    ):
        raise HTTPException(
            status_code=409,
            detail={"code": "SYSTEM_META_SIGNATURE_INVALID", "message": "_SYSTEM_META 签名校验失败"},
        )
    if manifest.get("schema_version") != META_SCHEMA or manifest.get("factory_id") != factory_id:
        raise HTTPException(
            status_code=409,
            detail={
                "code": "SYSTEM_META_SCOPE_MISMATCH",
                "message": "签名导出文件与当前厂区 scope 不一致",
            },
        )
    rows = manifest.get("rows")
    if (
        not isinstance(rows, list)
        or len(rows) != manifest.get("task_count")
        or _payload_hash(rows) != manifest.get("signed_row_manifest_digest")
    ):
        raise HTTPException(
            status_code=409,
            detail={
                "code": "SYSTEM_META_ROW_MANIFEST_INVALID",
                "message": "签名 row manifest 数量或摘要无效",
            },
        )
    required_row_fields = {
        "source_row",
        "row_type",
        "machine_id",
        "order_id",
        "task_id",
        "stable_row_key",
        "split_key",
        "quantity_scope",
        "planned_quantity",
        "order_revision",
        "task_revision",
        "task_report_counter_at_export",
        "takeover_source_completed_quantity",
        "completed_at_export",
        "report_increment_total_at_export",
        "progress_adjustment_total_at_export",
        "report_ledger_digest",
    }
    if any(
        not isinstance(row, dict) or required_row_fields - set(row)
        for row in rows
    ):
        raise HTTPException(
            status_code=409,
            detail={
                "code": "SYSTEM_META_ROW_MANIFEST_INCOMPLETE",
                "message": "签名 row manifest 缺少必需身份或进度水位",
            },
        )
    audit = db.scalar(
        select(InjectionSchedulingExportAudit).where(
            InjectionSchedulingExportAudit.id == manifest.get("export_id"),
            InjectionSchedulingExportAudit.factory_id == factory_id,
        )
    )
    if (
        audit is None
        or audit.manifest_json != manifest_json
        or audit.manifest_sha256 != supplied_manifest_hash
        or audit.metadata_signature != supplied_signature
        or audit.signing_key_id != supplied_key_id
        or audit.plan_id != manifest.get("plan_id")
        or audit.plan_revision != manifest.get("plan_revision")
        or audit.profile_id != manifest.get("profile_id")
        or audit.profile_revision != manifest.get("profile_revision")
        or audit.signed_row_manifest_digest
        != manifest.get("signed_row_manifest_digest")
    ):
        raise HTTPException(
            status_code=409,
            detail={
                "code": "SYSTEM_META_EXPORT_AUDIT_INVALID",
                "message": "_SYSTEM_META 对应的 export audit 不存在或不一致",
            },
        )
    plan = db.scalar(
        select(InjectionSchedulingPlan).where(
            InjectionSchedulingPlan.id == manifest.get("plan_id"),
            InjectionSchedulingPlan.factory_id == factory_id,
        )
    )
    if plan is None:
        raise HTTPException(
            status_code=409,
            detail={
                "code": "SYSTEM_META_PLAN_MISSING",
                "message": "签名导出引用的计划已不存在",
            },
        )
    task_ids = {str(row["task_id"]) for row in rows}
    tasks = {
        item.id: item
        for item in db.scalars(
            select(InjectionSchedulingTask).where(
                InjectionSchedulingTask.factory_id == factory_id,
                InjectionSchedulingTask.plan_id == manifest.get("plan_id"),
                InjectionSchedulingTask.id.in_(task_ids),
            )
        ).all()
    }
    order_ids = {str(row["order_id"]) for row in rows}
    orders = {
        item.id: item
        for item in db.scalars(
            select(InjectionSchedulingOrder).where(
                InjectionSchedulingOrder.factory_id == factory_id,
                InjectionSchedulingOrder.id.in_(order_ids),
            )
        ).all()
    }
    for row in rows:
        task = tasks.get(str(row["task_id"]))
        order = orders.get(str(row["order_id"]))
        if (
            task is None
            or order is None
            or task.order_id != order.id
            or task.machine_id != row["machine_id"]
            or task.stable_row_key != row["stable_row_key"]
        ):
            raise HTTPException(
                status_code=409,
                detail={
                    "code": "SYSTEM_META_ROW_IDENTITY_INVALID",
                    "message": "签名导出引用的 task/order/machine/stable identity 不一致",
                },
            )
    return manifest
