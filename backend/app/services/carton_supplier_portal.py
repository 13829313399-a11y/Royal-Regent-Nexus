"""Supplier-scoped collaboration. All mutations use the existing carton factory lock."""
import hashlib
import json
import re
from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import PurePosixPath
from types import SimpleNamespace
from uuid import uuid4
from zipfile import ZipFile, BadZipFile
from xml.etree import ElementTree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from fastapi import HTTPException
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session
from app.services.carton_purchase_batches import group_purchase_documents, load_purchase_batch, purchase_batch
from app.models.carton_mark import CartonMarkAsset, CartonMarkDocument, CartonMarkTemplate
from app.services import carton_mark_assets as mark_assets
from app.models.carton_procurement import CartonOrder, CartonOrderLine, CartonSupplier, CartonPurchaseOrderIssue, CartonReceipt, CartonReceiptLine, CartonAuditEvent
from app.models.carton_supplier_portal import SupplierCommitment, SupplierShipment, SupplierShipmentLine, SupplierShipmentUnmatchedLine, SupplierAttachment
from app.schemas.carton_supplier_portal import CommitmentSave, BatchCommitmentSave, ShipmentCreate, ShipmentReceive, CancellationAcknowledge, ShipmentStatusChange, SampleReceiptLink, ShipmentLineLink, SupplierMarkTemplateOut, SupplierDocumentExport
from app.schemas.carton_procurement import CartonReceiptCreate, CartonReceiptLineCreate, CartonReceiptConfirmRequest
from app.services.auth import AuthContext, authorization_decision, has_permission_in_scope
from app.services import carton_positions as positions
from app.services import carton_supplier_notifications as shipment_notifications
from app.services import carton_supplier_receipt_link as receipt_link
from app.services.carton_procurement_imports import _dongkang_identity, _dimension_identity
from app.services.carton_material_identity import material_conflicts, dimension_identity
from app.services.carton_procurement import (CARTON_DEPARTMENTS, require_carton_factory, _lock_receipt_factory,
    _purchase_order_issues, _purchase_order_snapshot, _purchase_order_pending_change, get_order_lines,
    fulfilled_by_line, _pending_received_by_line, _refresh_order_statuses, _audit, now_text, _create_receipt, confirm_receipt)

OPEN_STATES = {"PENDING_SUPPLIER", "PARTIALLY_RECEIVED"}
VISIBLE_STATES = OPEN_STATES | {"COMPLETED", "CANCELLED"}
MAX_FILE = 10 * 1024 * 1024
TABLES = ("carton_supplier_members", "carton_supplier_commitments", "carton_supplier_shipments", "carton_supplier_shipment_lines", "carton_supplier_unmatched_lines", "carton_supplier_attachments")


def _sample_link_id(receipt_line_id):
    return f"CAE-SAMPLE-LINK-{receipt_line_id}"


def _dimension_unit(value):
    match = re.search(r"(inch|in|英寸|cm|厘米)\s*$", str(value or ""), re.IGNORECASE)
    if not match:
        return ""
    return "in" if match.group(1).casefold() in {"inch", "in", "英寸"} else "cm"


def internal_permission(user, factory, *permissions):
    factory = require_carton_factory(factory)
    for permission in permissions:
        if not any(has_permission_in_scope(user, permission, factory, dept) for dept in CARTON_DEPARTMENTS):
            raise HTTPException(403, "没有当前厂区的供应商协同管理权限")
    return factory


def fixed_supplier(db, factory):
    supplier = db.scalar(select(CartonSupplier).where(CartonSupplier.factory_id == factory,
        CartonSupplier.supplier_code == "HEYUAN-DONGKANG", CartonSupplier.status == "ACTIVE"))
    if not supplier:
        raise HTTPException(409, "当前厂区未配置有效的河源东康供应商")
    return supplier


def supplier_factories(db, user):
    """Dongkang service factories with issued formal or published special requests."""
    if not supplier_permission(user, "carton_supplier:read", "*"):
        return []
    suppliers = db.scalars(select(CartonSupplier).where(CartonSupplier.supplier_code == "HEYUAN-DONGKANG",
        CartonSupplier.status == "ACTIVE").order_by(CartonSupplier.factory_id)).all()
    from app.models.carton_special_requests import SpecialRequest
    from app.services import carton_mark_sharing as sharing
    def has_shared_sources(supplier):
        audiences = sharing.states(db, supplier.factory_id)
        shared_ids = [identifier for identifier, state in audiences.items()
            if state['sharing_mode'] == 'SUPPLIER' and state['shared_supplier_id'] == supplier.id]
        return bool(shared_ids and db.scalar(select(CartonMarkAsset.id).where(
            CartonMarkAsset.factory_id == supplier.factory_id, CartonMarkAsset.is_archived.is_(False),
            CartonMarkAsset.id.in_(shared_ids)).limit(1)))
    return [supplier for supplier in suppliers if supplier_permission(
        user, "carton_supplier:read", supplier.factory_id) and (db.scalar(
        select(CartonOrder.id).join(CartonPurchaseOrderIssue, and_(
            CartonPurchaseOrderIssue.order_id == CartonOrder.id,
            CartonPurchaseOrderIssue.factory_id == CartonOrder.factory_id)).where(
                CartonOrder.factory_id == supplier.factory_id,
                CartonOrder.supplier_id == supplier.id,
                CartonOrder.status.in_(VISIBLE_STATES)).limit(1)) or db.scalar(select(SpecialRequest.id).where(
                    SpecialRequest.factory_id == supplier.factory_id, SpecialRequest.supplier_id == supplier.id,
                    SpecialRequest.status != "SUBMITTED").limit(1)) or has_shared_sources(supplier) or db.scalar(select(CartonAuditEvent.id).where(
                        CartonAuditEvent.factory_id == supplier.factory_id,
                        or_(and_(CartonAuditEvent.entity_type == "carton_supplier_payables", CartonAuditEvent.event_type == "SUPPLIER_PAYABLES_ACTION"),
                            and_(CartonAuditEvent.entity_type == "carton_price_correction", CartonAuditEvent.event_type == "RECEIPT_PRICE_CORRECTION_ACTION"))).limit(1)))]


def supplier_permission(user, permission, factory):
    # This module uses the evaluated IAM decision even during the legacy rollout;
    # its dedicated global overrides never change authorization in other modules.
    return (authorization_decision(user, permission, "*", "*")[0]
        and authorization_decision(user, permission, factory, "*")[0])


def supplier_access(db, user, factory, permission="carton_supplier:read"):
    factory = require_carton_factory(factory)
    supplier = next((item for item in supplier_factories(db, user) if item.factory_id == factory), None)
    if supplier is None or not supplier_permission(user, permission, factory):
        raise HTTPException(403, "账号无权访问此厂区的供应商订单")
    return supplier


def fingerprint(payload):
    data = payload.model_dump(mode="json")
    # Preserve request identities of ordinary shipments/receipts submitted before late-document support.
    if data.get("registration_mode") == "SHIPMENT":
        data.pop("registration_mode")
    if data.get("new_delivery_confirmation") is False:
        data.pop("new_delivery_confirmation")
    if data.get("remaining_action") == "CLOSE":
        data.pop("remaining_action")  # Preserve pre-batch receipt request fingerprints.
    for line in data.get("lines", []):
        if line.get("workshop_id") == "":
            line.pop("workshop_id")  # Preserve historical receipt retry fingerprints.
        if line.get("close_remaining") is False:
            line.pop("close_remaining", None)
        if line.get("replenishment_issue_id") is None:
            line.pop("replenishment_issue_id", None)  # Keep pre-bridge receipt retry identities.
        if line.get("unit_price", "provided") is None:
            line.pop("unit_price")
        if line.get("free_confirmed") is False:
            line.pop("free_confirmed")
        if line.get("price_action") == "MATCH":
            line.pop("price_action")
        if line.get("price_reason") == "":
            line.pop("price_reason")
    return hashlib.sha256(json.dumps(data, sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()


def latest_issue(db, order):
    return next((issue for issue in _purchase_order_issues(db, order.id)
        if not _purchase_order_snapshot(issue).get("replenishment")), None)


def acceptance_summary(db, order, lines, *, projection=None):
    """Internal read projection of supplier confirmation for the current ordinary issue."""
    from app.services.carton_supplier_changes import response_state
    issue = next((issue for issue in projection["issues"] if not _purchase_order_snapshot(issue).get("replenishment")), None) if projection is not None else latest_issue(db, order)
    positive_ids = [line.id for line in lines if line.required_quantity > 0]
    result = dict(status="NOT_ISSUED", label="尚未发送供应商",
        issue_id=issue.id if issue else "", document_no=issue.document_no if issue else "",
        total_line_count=len(positive_ids), accepted_line_count=0, accepted_at="", delivery_differences=[])
    change_state = (projection["supplier_change_states"].get(issue.id, {}) if projection is not None and "supplier_change_states" in projection
        else response_state(db, order.factory_id, "REGULAR", issue.id)) if issue else {}
    if order.status == "CANCELLED":
        notice = cancellation_notice(db, order, issue)
        label = "未下单已取消，无需通知供应商" if not issue else "取消通知已收到" if notice and notice["acknowledged"] else "取消通知待供应商确认" if notice and notice["issued"] else "订单已取消，变更待生成"
        result.update(status="CANCELLED", label=label)
    elif not issue or order.status not in VISIBLE_STATES:
        pass
    elif _purchase_order_pending_change(order, lines, issue)[0] != "NONE":
        result.update(status="PENDING_CHANGE", label="变更待生成")
    elif change_state.get("status") == "DISPUTED":
        result.update(status="PENDING", label="供应商有异议，暂停执行")
    elif not positive_ids:
        result.update(status="NOT_REQUIRED", label="无需接单")
    else:
        commitments = [row for row in projection["commitments"] if row.factory_id == order.factory_id and row.issue_id == issue.id and row.order_line_id in positive_ids] if projection is not None else list(db.scalars(select(SupplierCommitment).where(
            SupplierCommitment.factory_id == order.factory_id,
            SupplierCommitment.issue_id == issue.id,
            SupplierCommitment.order_line_id.in_(positive_ids))).all())
        accepted = len(commitments)
        status, label = ("ACCEPTED", "供应商已接单") if accepted == len(positive_ids) else (
            ("PARTIAL", "供应商部分接单") if accepted else ("PENDING", "供应商待接单"))
        result.update(status=status, label=label, accepted_line_count=accepted,
            accepted_at=max((row.accepted_at for row in commitments), default=""))
        papers = {line.id: line for line in lines}
        for commitment in sorted(commitments, key=lambda row: (papers[row.order_line_id].line_no, row.order_line_id)):
            try:
                days = (date.fromisoformat(commitment.promised_date) - date.fromisoformat(order.due_date)).days
            except (TypeError, ValueError):
                continue
            if days:
                result["delivery_differences"].append(dict(
                    order_line_id=commitment.order_line_id, packaging_type=papers[commitment.order_line_id].packaging_type,
                    planned_date=order.due_date, promised_date=commitment.promised_date, difference_days=days))
    return result


def acceptance_order_ids(db, orders, status):
    """Use the canonical current-issue projection before ledger counting/paging.

    Only three batched evidence reads are needed; a replenishment issue or an old
    commitment must never make the current ordinary purchase look accepted.
    """
    from collections import defaultdict
    result = []
    for start in range(0, len(orders), 400):
        batch = orders[start:start + 400]
        identifiers = [order.id for order in batch]
        lines, issues, commitments = defaultdict(list), defaultdict(list), defaultdict(list)
        for row in db.scalars(select(CartonOrderLine).where(CartonOrderLine.order_id.in_(identifiers))):
            lines[row.order_id].append(row)
        for row in db.scalars(select(CartonPurchaseOrderIssue).where(
                CartonPurchaseOrderIssue.order_id.in_(identifiers)).order_by(CartonPurchaseOrderIssue.issue_sequence.desc())):
            issues[row.order_id].append(row)
        line_ids = [line.id for own in lines.values() for line in own]
        for row in db.scalars(select(SupplierCommitment).where(SupplierCommitment.order_line_id.in_(line_ids))):
            commitments[row.order_line_id].append(row)
        from app.services.carton_supplier_changes import response_states
        change_states = {factory: response_states(db, factory, "REGULAR", [issue.id for own in issues.values() for issue in own if issue.factory_id == factory])
            for factory in {order.factory_id for order in batch}}
        for order in batch:
            own = lines[order.id]
            projection = {"issues": issues[order.id], "commitments": [row for line in own for row in commitments[line.id]],
                "supplier_change_states": change_states[order.factory_id]}
            if acceptance_summary(db, order, own, projection=projection)["status"] == status:
                result.append(order.id)
    return result


def supplier_order(db, order_id, factory, supplier_id):
    order = db.scalar(select(CartonOrder).where(CartonOrder.id == order_id, CartonOrder.factory_id == factory,
        CartonOrder.supplier_id == supplier_id, CartonOrder.deleted_at.is_(None), CartonOrder.status.in_(VISIBLE_STATES)))
    if order is None or latest_issue(db, order) is None:
        raise HTTPException(404, "未找到可访问的已下单订单")
    return order


def _mark_identity(customer_name, po, item, contract_number):
    return tuple(" ".join(value.strip().split()).casefold() for value in
        (customer_name, po, item, contract_number))


def _supplier_mark_orders(db, user, factory):
    """Current issued identities; unissued orders still take part in ambiguity checks."""
    supplier = supplier_access(db, user, factory)
    projections, eligible = [], {}
    for order in mark_assets.available_orders(db, factory):
        issue = latest_issue(db, order)
        header = _purchase_order_snapshot(issue).get("order", {}) if issue else {}
        if not isinstance(header, dict):
            header = {}
        # Unissued orders still participate in detecting cross-customer collisions.
        projections.append(SimpleNamespace(id=order.id, order_no=order.order_no,
            contract_no=str(header.get("contract_no", order.contract_no)),
            customer_code=str(header.get("customer_code", order.customer_code)),
            customer_name=str(header.get("customer_name", order.customer_name)),
            item_no=str(header.get("item_no", order.item_no))))
        if (issue and order.supplier_id == supplier.id
                and header.get("supplier_id") == supplier.id
                and order.status in OPEN_STATES | {"COMPLETED"}):
            eligible[order.id] = dict(id=order.id, issue_id=issue.id,
                customer_name=str(header.get("customer_name") or ""),
                contract_no=str(header.get("contract_no") or ""),
                customer_po=str(header.get("customer_po") or ""),
                item_no=str(header.get("item_no") or ""))
    return supplier, projections, eligible


def _supplier_mark_assets(db, user, factory, *, asset_ids=None):
    """Share source files by contract using issued headers, with factory-wide ambiguity checks."""
    from app.schemas.carton_mark import CartonMarkAssetWorkflow
    supplier, projections, eligible = _supplier_mark_orders(db, user, factory)
    from app.services import carton_mark_sharing as sharing
    audiences = sharing.states(db, factory)
    result = []
    query = select(CartonMarkAsset).where(CartonMarkAsset.factory_id == factory, CartonMarkAsset.is_archived.is_(False))
    if asset_ids is not None: query = query.where(CartonMarkAsset.id.in_(asset_ids))
    for asset in db.scalars(query.order_by(CartonMarkAsset.created_at.desc(), CartonMarkAsset.id.desc())):
        audience = audiences.get(asset.id, {})
        if audience.get('sharing_mode') == 'INTERNAL':
            continue
        explicitly_shared = audience.get('sharing_mode') == 'SUPPLIER'
        if explicitly_shared and audience.get('shared_supplier_id') != supplier.id:
            continue
        projected = mark_assets.asset_out(asset, projections, CartonMarkAssetWorkflow(), {})
        orders = [eligible[o.id] for o in projected.orders if o.id in eligible]
        from app.models.carton_special_requests import SpecialRequest, SpecialVersion
        special_visible = any((row := db.get(SpecialRequest, link['order_id'])) and row.factory_id == factory
            and row.supplier_id == supplier.id and row.status not in {'SUBMITTED', 'CANCELLED'}
            and not db.scalar(select(SpecialVersion.id).where(SpecialVersion.request_id == row.id, SpecialVersion.action == 'DELETED').limit(1))
            for link in audience.get('special_orders', []))
        if explicitly_shared or (projected.binding_status == 'BOUND' and orders) or special_visible:
            result.append((asset, orders))
    return result


def supplier_mark_assets(db, user, factory, *, asset_ids=None):
    from app.services import carton_mark_asset_workflow as lifecycle
    from app.schemas.carton_mark import CartonMarkAssetWorkflow
    visible = _supplier_mark_assets(db, user, factory, asset_ids=asset_ids)
    workflows = lifecycle.states(db, factory, [asset.id for asset, _ in visible])
    from app.services import carton_mark_sharing as sharing
    audiences = sharing.states(db, factory, [asset.id for asset, _ in visible])
    supplier, _, eligible = _supplier_mark_orders(db, user, factory)
    revisions = dict(db.execute(select(CartonOrder.id, CartonOrder.revision).where(
        CartonOrder.factory_id == factory, CartonOrder.id.in_(eligible))).all())
    def audience(asset, orders):
        state = dict(audiences.get(asset.id, {}))
        state['pending_bindings'] = [x for x in state.get('pending_bindings', [])
            if x['supplier_id'] == supplier.id and x['order_id'] in eligible]
        from app.models.carton_special_requests import SpecialRequest
        from app.models.carton_special_requests import SpecialVersion
        state['special_orders'] = [x for x in state.get('special_orders', []) if (row := db.get(SpecialRequest, x['order_id']))
            and row.factory_id == factory and row.supplier_id == supplier.id and row.status not in {'SUBMITTED', 'CANCELLED'}
            and not db.scalar(select(SpecialVersion.id).where(SpecialVersion.request_id == row.id, SpecialVersion.action == 'DELETED').limit(1))]
        from app.services.carton_mark_associations import projection
        # The public identity belongs to the issued purchase order. Only the
        # compare-and-save revision comes from the mutable internal order.
        state['associations'] = [dict(order_kind='REGULAR', order_id=order['id'],
            contract_no=order['contract_no'], customer_name=order['customer_name'], item_no=order['item_no'],
            expected_revision=revisions[order['id']], issue_id=order['issue_id'], line_id='')
            for order in orders if order['id'] in eligible]
        state['associations'] += projection(db, asset, [], state)
        return state
    result = []
    for asset, orders in visible:
        state = audience(asset, orders)
        result.append(dict(id=asset.id, revision=asset.revision, file_name=asset.file_name, kind=asset.kind,
            photo_group_id=asset.photo_group_id,
            size_bytes=asset.size_bytes, contract_number="、".join(sorted({order["contract_no"] for order in orders})),
            created_at=asset.created_at, orders=orders, binding_status='BOUND' if orders or state.get('special_orders') else 'UNBOUND', **state,
            workflow=workflows.get(asset.id, CartonMarkAssetWorkflow()).model_dump()))
    return result


def supplier_mark_asset_document(db, user, factory, asset_id):
    asset = next((asset for asset, _ in _supplier_mark_assets(db, user, factory)
        if asset.id == asset_id), None)
    if asset is None:
        raise HTTPException(404, "未找到可查看的箱唛资料")
    return asset


def _visible_supplier_mark_templates(db, user, factory):
    supplier = supplier_access(db, user, factory)
    orders = db.scalars(select(CartonOrder).where(
        CartonOrder.factory_id == factory, CartonOrder.supplier_id == supplier.id,
        CartonOrder.status.in_(OPEN_STATES | {"COMPLETED"}))).all()
    issued_identities = set()
    for order in orders:
        issue = latest_issue(db, order)
        if issue is None:
            continue
        header = _purchase_order_snapshot(issue).get("order", {})
        if not isinstance(header, dict) or header.get("supplier_id") != supplier.id:
            continue
        issued_identities.add(_mark_identity(str(header.get("customer_name") or ""),
            str(header.get("customer_po") or header.get("contract_no") or ""),
            str(header.get("item_no") or ""), str(header.get("contract_no") or "")))
    if not issued_identities:
        return []
    candidates = db.scalars(select(CartonMarkTemplate).where(
        CartonMarkTemplate.factory_id == factory)).all()
    latest_by_identity = {}
    for template in candidates:
        identity = _mark_identity(template.customer_name, template.po, template.item,
            template.contract_number)
        if identity not in issued_identities:
            continue
        previous = latest_by_identity.get(identity)
        if previous is None or (template.version, template.created_at, template.id) > (
            previous.version, previous.created_at, previous.id):
            latest_by_identity[identity] = template
    return sorted((template for template in latest_by_identity.values()
        if not template.is_archived and (template.check_status == "核对通过" or template.manual_released_at)),
        key=lambda template: (template.created_at, template.id), reverse=True)


def supplier_mark_templates(db, user, factory):
    templates = _visible_supplier_mark_templates(db, user, factory)
    if not templates:
        return []
    documents = {}
    for template_id, kind, filename in db.execute(select(
        CartonMarkDocument.template_id, CartonMarkDocument.kind, CartonMarkDocument.file_name).where(
        CartonMarkDocument.factory_id == factory,
        CartonMarkDocument.template_id.in_([template.id for template in templates]))):
        documents.setdefault(template_id, {})[kind] = filename
    return [SupplierMarkTemplateOut(id=template.id, customer_name=template.customer_name,
        po=template.po, item=template.item, contract_number=template.contract_number,
        version=template.version, check_status=template.check_status,
        manual_released=bool(template.manual_released_at),
        excel_file_name=documents[template.id]["source_excel"],
        pdf_file_name=documents[template.id]["print_pdf"], created_at=template.created_at)
        for template in templates if {"source_excel", "print_pdf"} <= documents.get(template.id, {}).keys()]


def supplier_mark_document(db, user, factory, template_id, kind):
    if kind not in {"source_excel", "print_pdf"}:
        raise HTTPException(422, "箱唛文档类型无效")
    if not any(template.id == template_id for template in _visible_supplier_mark_templates(db, user, factory)):
        raise HTTPException(404, "未找到可查看的箱唛资料")
    available_kinds = set(db.scalars(select(CartonMarkDocument.kind).where(
        CartonMarkDocument.factory_id == factory,
        CartonMarkDocument.template_id == template_id)).all())
    if not {"source_excel", "print_pdf"} <= available_kinds:
        raise HTTPException(404, "箱唛文档不完整")
    document = db.scalar(select(CartonMarkDocument).where(
        CartonMarkDocument.factory_id == factory,
        CartonMarkDocument.template_id == template_id, CartonMarkDocument.kind == kind))
    if document is None:
        raise HTTPException(404, "箱唛文档不存在")
    return document


def executable(db, order, issue_id, *, allow_dispute=False):
    if order.deleted_at:
        raise HTTPException(409, "订单已归档删除，不能继续接单或发货")
    if order.status not in OPEN_STATES:
        raise HTTPException(409, "订单当前状态不允许接单或发货")
    issue = latest_issue(db, order)
    if not issue or issue.id != issue_id:
        raise HTTPException(409, "采购版本已变更，请刷新后重新确认接单")
    if _purchase_order_pending_change(order, get_order_lines(db, order.id), issue)[0] != "NONE":
        raise HTTPException(409, "订单存在未生成变更，旧承诺暂停执行，请等待内部生成后重新接单")
    from app.services.carton_supplier_changes import response_state
    if not allow_dispute and response_state(db, order.factory_id, "REGULAR", issue.id)["status"] == "DISPUTED":
        raise HTTPException(409, "供应商已提出异议，请在变更待办核对整单后重新同意")
    return issue


def outstanding(db, line_ids):
    if not line_ids:
        return {}
    rows = db.execute(select(SupplierShipmentLine, SupplierShipment, CartonReceipt.status)
        .join(SupplierShipment, SupplierShipment.id == SupplierShipmentLine.shipment_id)
        .outerjoin(CartonReceipt, CartonReceipt.id == SupplierShipment.receipt_id)
        .where(or_(SupplierShipment.status == "SENT", and_(SupplierShipment.status == "RECEIVED", CartonReceipt.status == "REVERSED")),
            SupplierShipmentLine.order_line_id.in_(line_ids))).all()
    result = {}
    from app.services import carton_shipment_receiving as receiving
    projections = {}
    for paper, shipment, receipt_status in rows:
        if receipt_link.link_only(shipment):
            continue  # A document for already posted goods is not another in-transit quantity.
        if shipment.id not in projections:
            formal, unmatched = _shipment_sources(db, shipment)
            projections[shipment.id] = receiving.progress(db, shipment, formal + unmatched)[2]
        reserved = projections[shipment.id][paper.id]["reserved_quantity"]
        result[paper.order_line_id] = result.get(paper.order_line_id, Decimal(0)) + reserved
    return result


def _shipment_sources(db, shipment):
    formal = list(db.scalars(select(SupplierShipmentLine).where(
        SupplierShipmentLine.shipment_id == shipment.id).order_by(SupplierShipmentLine.id)).all())
    linked = {json.loads(line.snapshot_json).get("original_unmatched_line_id") for line in formal}
    unmatched = [line for line in db.scalars(select(SupplierShipmentUnmatchedLine).where(
        SupplierShipmentUnmatchedLine.shipment_id == shipment.id).order_by(SupplierShipmentUnmatchedLine.id)).all()
        if line.id not in linked]
    return formal, unmatched


def blocked_replacement_lines(db, factory, line_ids):
    from app.services.carton_replenishment_receipts import options_by_line, legacy_review_lines
    return set(options_by_line(db, factory, line_ids)) | legacy_review_lines(db, factory, line_ids)


def attachment_out(row):
    return {key: getattr(row, key) for key in ("id", "filename", "version", "size", "sha256", "created_at")}


def _snapshot_line_no(value, order_no):
    if isinstance(value, int) and not isinstance(value, bool):
        number = value
    elif isinstance(value, str) and re.fullmatch(r"[0-9]+", value.strip()):
        number = int(value.strip())
    else:
        raise HTTPException(409, f"订单 {order_no} 的采购快照纸品行号无效，请核对原始采购记录")
    if number <= 0:
        raise HTTPException(409, f"订单 {order_no} 的采购快照纸品行号必须为正整数，请核对原始采购记录")
    return number


def order_out(db, order, *, internal=False):
    issue = latest_issue(db, order)
    if not issue and not internal:
        return None
    snapshot = _purchase_order_snapshot(issue) if issue else _purchase_order_pending_change(order, get_order_lines(db, order.id), None)[3]
    header = snapshot.get("order", {})
    current_lines = get_order_lines(db, order.id)
    ids = [line.id for line in current_lines]
    fulfilled = fulfilled_by_line(db, ids)
    pending = _pending_received_by_line(db, ids)
    shipped = outstanding(db, ids)
    from app.services import carton_replenishment_supplier as replacement
    from app.services.carton_replenishment_receipts import legacy_review_lines
    blocked = legacy_review_lines(db, order.factory_id, ids)
    changed = _purchase_order_pending_change(order, current_lines, issue)[0] != "NONE"
    result = {key: header.get(key, "") for key in ("order_no", "customer_name", "contract_no", "customer_po", "item_no", "product_name")}
    result.update(id=order.id, status=order.status, revision=order.revision,
        issue_id=issue.id if issue else "", document_no=issue.document_no if issue else "尚未生成",
        order_date=header.get("order_date") or order.order_date,
        planned_date=snapshot.get("after_due_date", ""), awaiting_issue=changed, lines=[])
    from app.services.carton_material_correction import records, blocked_lines
    result["material_corrections"] = records(db, order.factory_id, order.id)
    material_holds = blocked_lines(db, order.factory_id, ids)
    result["cancellation_notice"] = cancellation_notice(db, order, issue)
    from app.services.carton_supplier_changes import response_state
    disputed = bool(issue and response_state(db, order.factory_id, "REGULAR", issue.id)["status"] == "DISPUTED")
    if internal:
        from app.services.carton_order_split import plans
        result["customer_code"] = order.customer_code
        result["split_records"] = plans(db, order.factory_id, order.id)
    for item in snapshot.get("lines", []):
        line_id = str(item["id"])
        required = Decimal(str(item["after_required_quantity"]))
        commitment = db.get(SupplierCommitment, line_id)
        accepted = bool(issue and commitment and commitment.issue_id == issue.id and not changed and not disputed)
        line = {key: item.get(key, "") for key in ("id", "line_no", "packaging_type", "paper_quality", "specification", "dimension_unit", "unit")}
        line_no = _snapshot_line_no(item.get("line_no"), order.order_no)
        line.update(line_no=line_no, child_no=f"{order.order_no}/{line_no:02d}", required_quantity=str(required),
            received_quantity=str(fulfilled.get(line_id, 0)), in_transit_quantity=str(shipped.get(line_id, 0)),
            remaining_to_ship=str(replacement.ordinary_capacity(db, order, line_id)),
            shipping_blocked_reason="供应商有异议，暂停执行，请沿变更待办核实" if disputed else "纸品资料更正待双方确认，暂停发货和正式入库" if line_id in material_holds else "此纸品有补单待核对，请联系内部仓库按原补单流程处理" if line_id in blocked else "",
            accepted=accepted, commitment_revision=commitment.revision if commitment else 0,
            promised_date=commitment.promised_date if commitment else "")
        if internal:
            line.update(unit_price=item.get("unit_price", "0"), currency=item.get("currency", "CNY"))
        result["lines"].append(line)
    result["replenishment_sources"] = replacement.sources(db, order)
    result["attachments"] = [attachment_out(row) for row in db.scalars(select(SupplierAttachment).where(
        SupplierAttachment.order_id == order.id, SupplierAttachment.factory_id == order.factory_id).order_by(SupplierAttachment.created_at.desc())).all()]
    return result


def cancellation_notice(db, order, issue=None):
    if order.status != "CANCELLED":
        return None
    issue = issue or latest_issue(db, order)
    if not issue:
        return None
    issued = bool(issue and _purchase_order_snapshot(issue).get("cancellation"))
    event = db.scalar(select(CartonAuditEvent).where(
        CartonAuditEvent.factory_id == order.factory_id,
        CartonAuditEvent.entity_type == "carton_order", CartonAuditEvent.entity_id == order.id,
        CartonAuditEvent.event_type == "SUPPLIER_CANCELLATION_ACKNOWLEDGED")
        .order_by(CartonAuditEvent.sequence.desc()).limit(1)) if issued else None
    saved = json.loads(event.detail_json) if event else {}
    acknowledged = bool(event and saved.get("issue_id") == issue.id
                        and (saved.get("order_revision") == order.revision or (
                            order.deleted_at and saved.get("order_revision") == issue.source_order_revision)))
    return {"issue_id": issue.id if issued else "", "order_revision": order.revision,
            "issued": issued, "acknowledged": acknowledged,
            "acknowledged_at": event.created_at if acknowledged else ""}


def acknowledge_cancellation(db, user, order_id, payload: CancellationAcknowledge, *, commit=True):
    _lock_receipt_factory(db, payload.factory_id)
    supplier = supplier_access(db, user, payload.factory_id, "carton_supplier:approve")
    order = supplier_order(db, order_id, payload.factory_id, supplier.id)
    issue = latest_issue(db, order)
    if order.status != "CANCELLED" or not issue or not _purchase_order_snapshot(issue).get("cancellation"):
        raise HTTPException(409, "订单尚无当前取消变更单，请刷新并核对")
    if (issue.id != payload.issue_id or order.revision != payload.expected_order_revision
            or issue.source_order_revision != order.revision):
        raise HTTPException(409, "取消通知版本已变更，请刷新后确认")
    if not cancellation_notice(db, order, issue)["acknowledged"]:
        _audit(db, user, payload.factory_id, "SUPPLIER_CANCELLATION_ACKNOWLEDGED", "carton_order", order.id,
               {"issue_id": issue.id, "order_revision": order.revision, "document_no": issue.document_no})
    db.commit() if commit else db.flush()
    return order_out(db, order)


def ensure_shipment_customers(db, user, shipment):
    from app.services.carton_customer_assignment import ensure_customer_operation
    formal, unmatched = _shipment_sources(db, shipment)
    codes = set(db.scalars(select(CartonOrder.customer_code).join(CartonOrderLine, CartonOrderLine.order_id == CartonOrder.id)
        .where(CartonOrder.factory_id == shipment.factory_id, CartonOrderLine.id.in_([line.order_line_id for line in formal]))))
    codes.update(json.loads(line.snapshot_json).get("customer_code", "") for line in unmatched)
    for code in codes:
        ensure_customer_operation(db, user, shipment.factory_id, code)


def visible_shipment(db, user, shipment):
    try:
        ensure_shipment_customers(db, user, shipment)
        return True
    except HTTPException as exc:
        if exc.status_code != 403:
            raise
        return False


def pending_shipments(db, user, factory, *, limit=3):
    """The dashboard counts original notes across all authorized paper sources."""
    from app.services.carton_unified_deliveries import pending
    return pending(db, user, factory, limit=limit)


def workspace(db, user, factory, *, internal=False, responsibility_scope="OWN"):
    if internal:
        internal_permission(user, factory, "carton_procurement:read")
        supplier_id = fixed_supplier(db, factory).id
    else:
        supplier_id = supplier_access(db, user, factory).id
    orders = db.scalars(select(CartonOrder).where(CartonOrder.factory_id == factory, CartonOrder.supplier_id == supplier_id,
        CartonOrder.deleted_at.is_(None),
        CartonOrder.status.in_(VISIBLE_STATES | {"DRAFT", "CONFIRMED"} if internal else VISIBLE_STATES)).order_by(CartonOrder.due_date, CartonOrder.order_no)).all()
    if internal and responsibility_scope == "OWN":
        from app.services.carton_customer_assignment import customer_predicate
        visible_order_ids = set(db.scalars(select(CartonOrder.id).where(CartonOrder.factory_id == factory,
            customer_predicate(user, factory, CartonOrder.customer_code))))
        orders = [order for order in orders if order.id in visible_order_ids]
    result = [order_out(db, row, internal=internal) for row in orders]
    shipments = db.scalars(select(SupplierShipment).where(SupplierShipment.factory_id == factory,
        SupplierShipment.supplier_id == supplier_id).order_by(SupplierShipment.created_at.desc())).all()
    if internal and responsibility_scope == "OWN":
        shipments = [row for row in shipments if visible_shipment(db, user, row)]
    return {"factory_id": factory, "supplier_name": db.get(CartonSupplier, supplier_id).supplier_name,
        "orders": [row for row in result if row], "shipments": [dict(shipment_out(db, row, internal=internal), **({"can_operate": visible_shipment(db, user, row)} if internal else {})) for row in shipments]}


def _closed_source_linkable(shipment, source, history, progress):
    from app.services import carton_shipment_receiving as receiving
    return bool(shipment.status == "NOT_RECEIVED" and not shipment.receipt_id and history
        and progress[source.id]["closed_quantity"] > 0
        and all(not batch.get("legacy_unsequenced")
            and receiving.entry(batch, source)
            and receiving.amount(receiving.entry(batch, source), "received_quantity") == 0 for batch in history))


def shipment_out(db, row, *, internal=False):
    result = {key: getattr(row, key) for key in ("id", "delivery_note_no", "delivery_date", "status", "revision", "created_at", "confirmed_at")}
    receipt = db.get(CartonReceipt, row.receipt_id) if row.receipt_id else None
    correction = bool(receipt and receipt.status == "REVERSED")
    result["requires_correction"] = correction
    if correction:
        result["status"] = "RECEIPT_REVERSED"
    result["lines"] = []
    result["source_filename"] = ""
    result["source_sha256"] = ""
    formal, unmatched = _shipment_sources(db, row)
    from app.services import carton_shipment_receiving as receiving
    batches, correction_batch, progress = receiving.progress(db, row, formal + unmatched)
    correction = correction_batch is not None
    result["requires_correction"] = correction
    if correction:
        result["status"] = "RECEIPT_REVERSED"
        receipt = correction_batch["receipt"]
    for line in formal:
        snapshot = json.loads(line.snapshot_json)
        source_file = snapshot.get("source_file") or {}
        result["source_filename"] = result["source_filename"] or source_file.get("filename", "")
        result["source_sha256"] = result["source_sha256"] or source_file.get("sha256", "")
        item = {key: snapshot.get(key, "") for key in ("order_no", "contract_no", "customer_po", "item_no", "customer_name", "child_no", "packaging_type", "paper_quality", "specification", "unit")}
        item.update(id=line.id, order_line_id=line.order_line_id, quantity=str(line.quantity), source_type="FORMAL_ORDER")
        item.update(replenishment_issue_id=snapshot.get("replenishment_issue_id"),
            replenishment_document_no=snapshot.get("replenishment_document_no"), responsibility=snapshot.get("responsibility"))
        from app.services.carton_material_correction import latest_approved, receipt_specification, blocked_lines
        item["material_correction_pending"] = bool(blocked_lines(db, row.factory_id, [line.order_line_id]))
        if (row.status == 'SENT' or correction) and latest_approved(db, row.factory_id, line.order_line_id):
            current_paper = db.get(CartonOrderLine, line.order_line_id)
            item.update(packaging_type=current_paper.packaging_type, paper_quality=current_paper.paper_quality,
                specification=receipt_specification(db, current_paper), dimension_unit=current_paper.dimension_unit,
                material_correction_note='已按双方确认的资料更正；原送货资料保留')
        if internal:
            item.update(unit_price=snapshot.get("unit_price", "0"),
                delivery_unit_price=snapshot.get("delivery_unit_price"), currency=snapshot.get("currency", "CNY"))
            if snapshot.get("original_unmatched_line_id"):
                item["original_unmatched_line_id"] = snapshot["original_unmatched_line_id"]
            original = snapshot.get("original_source_material")
            if original:
                item["source_material"] = {key: original.get(key) for key in ("packaging_type", "paper_quality", "specification", "source_sheet", "source_row")}
                if original.get("packaging_type_explicit") is False:
                    item["source_material"]["packaging_type"] = ""
        result["lines"].append(item)
    for line in unmatched:
        snapshot = json.loads(line.snapshot_json)
        source_file = snapshot.get("source_file") or {}
        result["source_filename"] = result["source_filename"] or source_file.get("filename", "")
        result["source_sha256"] = result["source_sha256"] or source_file.get("sha256", "")
        item = {key: snapshot.get(key, "") for key in ("order_no", "contract_no", "customer_po", "item_no", "customer_name", "child_no", "packaging_type", "paper_quality", "specification", "unit")}
        item.update(id=line.id, order_line_id=None, quantity=str(line.quantity), source_type="AD_HOC_REVIEW")
        if internal:
            item.update(unit_price=str(snapshot.get("unit_price") or "0"), delivery_unit_price=snapshot.get("unit_price"),
                free_confirmed=snapshot.get("free_confirmed") is True or snapshot.get("pricing_kind") == "FREE", currency="CNY",
                closed_source_linkable=_closed_source_linkable(row, line, batches, progress))
        classified = next((receiving.entry(batch, line) for batch in reversed(batches)
            if batch.get("receipt") and batch["receipt"].status == "POSTED"
            and receiving.entry(batch, line).get("no_order_decision") in {"SAMPLE", "MISC", "GIFT"}), None)
        if classified:
            item["business_kind"] = {"SAMPLE": "SAMPLE", "MISC": "MISC", "GIFT": "UNKNOWN"}[classified["no_order_decision"]]
            if internal:
                item["receipt_classification"] = {key: classified.get(key, "") for key in
                    ("no_order_decision", "customer_code", "requested_by", "sample_purpose", "workshop_id")}
        result["lines"].append(item)
    accepted = json.loads(row.acceptance_json or "{}")
    for item in result["lines"]:
        item.update({key: str(value) for key, value in progress[item["id"]].items() if key != "reserved_quantity"})
    result["partially_received"] = row.status == "SENT" and any(item["received_quantity"] > 0 for item in progress.values())
    result["requires_receipt_link"] = receipt_link.link_only(row)
    result["linked_existing_receipt"] = bool(accepted.get("linked_existing_receipt"))
    result["acceptance_date"] = accepted.get("acceptance_date")
    # Supplier sees actual quantities/discrepancies, never internal cost or warehouse IDs.
    result["acceptance_lines"] = [{key: item.get(key) for key in ("shipment_line_id", "received_quantity", "damaged_quantity", "rejected_quantity", "unusable_quantity", "difference_reason", "no_order_decision")} for item in accepted.get("lines", [])]
    if batches:
        result["acceptance_lines"] = []
        for paper in formal + unmatched:
            entries = [receiving.entry(batch, paper) for batch in batches if receiving.allocated(batch, paper) > 0
                and (not batch["receipt"] or batch["receipt"].status == "POSTED")]
            if not entries:
                continue
            result["acceptance_lines"].append({"shipment_line_id": paper.id,
                **{key: str(progress[paper.id][key]) for key in receiving.QUANTITIES},
                "difference_reason": "；".join(dict.fromkeys(item.get("difference_reason", "") for item in entries if item.get("difference_reason"))),
                "no_order_decision": next((item.get("no_order_decision") for item in reversed(entries) if item.get("no_order_decision")), "")})
    last_status_event = db.scalar(select(CartonAuditEvent).where(CartonAuditEvent.factory_id == row.factory_id,
        CartonAuditEvent.entity_type == "supplier_shipment", CartonAuditEvent.entity_id == row.id,
        CartonAuditEvent.event_type.in_({"SUPPLIER_SHIPMENT_DEFERRED", "SUPPLIER_SHIPMENT_NON_ARRIVAL_REOPENED"}))
        .order_by(CartonAuditEvent.sequence.desc()).limit(1))
    pending_detail = json.loads(last_status_event.detail_json) if last_status_event else {}
    result["pending_reason"] = pending_detail.get("reason", "") if row.status == "SENT" and pending_detail.get("revision") == row.revision else ""
    result["acceptance_history"] = []
    events = db.scalars(select(CartonAuditEvent).where(CartonAuditEvent.factory_id == row.factory_id,
        CartonAuditEvent.entity_type == "supplier_shipment", CartonAuditEvent.entity_id == row.id,
        CartonAuditEvent.event_type.in_({"SUPPLIER_SHIPMENT_RECEIVED", "SUPPLIER_SHIPMENT_NOT_RECEIVED", receipt_link.EVENT}))
        .order_by(CartonAuditEvent.sequence)).all()
    for event in events:
        detail = json.loads(event.detail_json)
        prior = db.get(CartonReceipt, detail.get("receipt_id")) if detail.get("receipt_id") else None
        acceptance = detail.get("acceptance") or {}
        result["acceptance_history"].append({**({"receipt_id": prior.id} if internal and prior else {}), "confirmed_at": event.created_at,
            "acceptance_date": acceptance.get("acceptance_date"), "status": prior.status if prior else ("NOT_RECEIVED" if event.event_type == "SUPPLIER_SHIPMENT_NOT_RECEIVED" else "RECEIVED"),
            "lines": [{key: item.get(key) for key in ("shipment_line_id", "received_quantity", "damaged_quantity", "rejected_quantity", "unusable_quantity", "difference_reason", "no_order_decision")}
                for item in acceptance.get("lines", [])]})
    if correction and not any(batch["receipt"] and batch["receipt"].status == "POSTED" for batch in batches):
        result["acceptance_date"] = None
        result["acceptance_lines"] = []  # Old acceptance is history, not current received quantity.
    if internal:
        result["legacy_correction_lines"] = [item for item in (correction_batch["acceptance"].get("lines", []) if correction_batch else [])
            if item.get("no_order_decision") in {"SAMPLE", "MISC", "GIFT", "UNORDERED"} and Decimal(str(item.get("received_quantity") or 0)) > 0]
        result["receipt_id"] = receipt.id if receipt else row.receipt_id
        result["receipt_status"] = receipt.status if receipt else None
        receipt_lines = db.scalars(select(CartonReceiptLine).where(
            CartonReceiptLine.receipt_id.in_([batch["receipt_id"] for batch in batches if batch["receipt"] and batch["receipt"].status == "POSTED"]),
            CartonReceiptLine.source_type == "AD_HOC").order_by(CartonReceiptLine.line_no)).all()
        result["sample_receipts"] = []
        for receipt_line in receipt_lines:
            link = db.scalar(select(CartonAuditEvent).where(
                CartonAuditEvent.id == _sample_link_id(receipt_line.id),
                CartonAuditEvent.factory_id == row.factory_id))
            result["sample_receipts"].append({
                "receipt_line_id": receipt_line.id, "customer_code": receipt_line.customer_code,
                "customer_name": receipt_line.customer_name, "contract_no": receipt_line.contract_no,
                "item_no": receipt_line.item_no, "packaging_type": receipt_line.packaging_type,
                "paper_quality": receipt_line.paper_quality, "specification": receipt_line.specification,
                "unit": receipt_line.unit, "quantity": str(receipt_line.effective_quantity),
                "linked_order_line_id": link.entity_id if link else "",
            })
    return result


def documents(db, user, factory):
    """Issued purchase snapshots and supplier-authored notes only, within one authorized factory."""
    supplier = supplier_access(db, user, factory)
    orders = db.scalars(select(CartonOrder).where(
        CartonOrder.factory_id == factory, CartonOrder.supplier_id == supplier.id,
        CartonOrder.status.in_(VISIBLE_STATES))).all()
    order_ids = [order.id for order in orders]
    papers, commitments, all_issues = defaultdict(list), defaultdict(list), defaultdict(list)
    for start in range(0, len(order_ids), 400):
        identifiers = order_ids[start:start + 400]
        for paper in db.scalars(select(CartonOrderLine).where(CartonOrderLine.factory_id == factory, CartonOrderLine.order_id.in_(identifiers))):
            papers[paper.order_id].append(paper)
        for commitment, order_id in db.execute(select(SupplierCommitment, CartonOrderLine.order_id).join(
                CartonOrderLine, CartonOrderLine.id == SupplierCommitment.order_line_id).where(
                SupplierCommitment.factory_id == factory, CartonOrderLine.factory_id == factory,
                CartonOrderLine.order_id.in_(identifiers))):
            commitments[order_id].append(commitment)
        for issue in db.scalars(select(CartonPurchaseOrderIssue).where(
                CartonPurchaseOrderIssue.factory_id == factory, CartonPurchaseOrderIssue.order_id.in_(identifiers))
                .order_by(CartonPurchaseOrderIssue.issue_sequence.desc())):
            all_issues[issue.order_id].append(issue)
    acceptances = {order.id: acceptance_summary(db, order, papers[order.id], projection={
        "issues": all_issues[order.id], "commitments": commitments[order.id]}) for order in orders}
    result = []
    batches = {}
    if order_ids:
        issues = db.scalars(select(CartonPurchaseOrderIssue).where(
            CartonPurchaseOrderIssue.factory_id == factory,
            CartonPurchaseOrderIssue.order_id.in_(order_ids),
            CartonPurchaseOrderIssue.document_type != "LEGACY_BASELINE")
            .order_by(CartonPurchaseOrderIssue.generated_at.desc())).all()
        for issue in issues:
            snapshot = _purchase_order_snapshot(issue)
            batches[issue.id] = purchase_batch(issue)
            header = snapshot.get("order", {})
            acceptance = dict(acceptances[issue.order_id])
            if snapshot.get("replenishment"):
                from app.services import carton_replenishment_supplier as replacement
                source = next(item for item in replacement.sources(db, db.get(CartonOrder, issue.order_id)) if item["issue_id"] == issue.id)
                acceptance = replacement.acceptance(source)
            elif acceptance["issue_id"] != issue.id:
                acceptance.update(status="HISTORICAL", label="旧版本，接单需核实", accepted_line_count=0, accepted_at="", delivery_differences=[])
            acceptance.update(issue_id=issue.id, document_no=issue.document_no)
            result.append({
                "id": issue.id, "kind": "PURCHASE", "factory_id": factory,
                "document_no": issue.document_no, "document_type": issue.document_type,
                "date": issue.generated_at[:10], "created_at": issue.generated_at,
                "status": "已生成", "replenishment": bool(snapshot.get("replenishment")),
                "is_cancellation": bool(snapshot.get("cancellation")),
                "supplier_acceptance": acceptance,
                "orders": [{
                    "order_kind": "REGULAR",
                    "order_no": issue.order_no, "customer_name": header.get("customer_name") or "",
                    "contract_no": header.get("contract_no") or "", "customer_po": header.get("customer_po") or "",
                    "item_no": header.get("item_no") or "", "product_name": header.get("product_name") or "",
                    "order_date": header.get("order_date") or "",
                    "planned_date": snapshot.get("after_due_date") or "",
                }],
                "lines": [{
                    "order_kind": "REGULAR",
                    "order_no": issue.order_no, "child_no": f"{issue.order_no}/{int(line.get('line_no', 0)):02d}",
                    "packaging_type": line.get("packaging_type", ""),
                    "paper_quality": line.get("paper_quality", ""),
                    "specification": line.get("specification", ""),
                    "unit": line.get("unit", ""),
                    "before_quantity": str(line.get("before_required_quantity", "0")),
                    "change_quantity": str(line.get("required_quantity_delta", "0")),
                    "quantity": str(line.get("after_required_quantity", "0")),
                } for line in snapshot.get("lines", [])],
            })
            if snapshot.get("replenishment"):
                result[-1]["lines"] = [dict(line, order_kind="REGULAR", order_no=issue.order_no,
                    quantity=line["required_quantity"], change_quantity=line["required_quantity"], before_quantity="0") for line in source["lines"]]
                result[-1]["responsibility"] = source["responsibility"]
    result = group_purchase_documents(result, batches)
    from app.services.carton_special_purchase import supplier_documents
    result.extend(supplier_documents(db, user, factory))
    from app.services.carton_unified_deliveries import workspace as unified_workspace
    for group in unified_workspace(db, user, factory, external=True)["items"]:
        visible = group["formal"] or {}
        special_by_id = {row["id"]: row for row in group["special"]}
        def business_kind(line):
            if line["source_kind"] == "FORMAL":
                if line.get("source_type") == "AD_HOC_REVIEW":
                    return "UNKNOWN"
                return "REGULAR" if line.get("source_type") == "FORMAL_ORDER" or line.get("order_line_id") else "UNKNOWN"
            source = special_by_id.get(line.get("special_request_id"))
            if not source or source.get("origin") == "SUPPLIER_EXTRA" and not any(shipment.get("receipts") for shipment in source["shipments"]):
                return "UNKNOWN"
            return {"SAMPLE": "SAMPLE", "WORKSHOP_MISC": "MISC"}.get(source.get("kind"), "UNKNOWN")
        orders_by_no = {}
        for line in group["lines"]:
            if not line.get("order_no"):
                continue
            orders_by_no.setdefault(line["order_no"], {
                "order_kind": business_kind(line),
                "order_no": line["order_no"], "customer_name": line.get("customer_name", ""),
                "contract_no": line.get("contract_no", ""), "customer_po": line.get("customer_po", ""),
                "item_no": line.get("item_no", ""), "product_name": "", "order_date": "", "planned_date": "",
            })
        result.append({
            "id": group["id"].removeprefix("FORMAL:"), "kind": "DELIVERY", "factory_id": factory,
            "document_no": group["delivery_note_no"], "document_type": "DELIVERY",
            "date": group["delivery_date"], "created_at": group["created_at"],
            "status": group["status"], "replenishment": False,
            "source_filename": visible.get("source_filename", ""), "source_sha256": visible.get("source_sha256", ""),
            "orders": list(orders_by_no.values()),
            "unmatched_line_count": sum(line.get("source_type") == "AD_HOC_REVIEW" for line in group["lines"]),
            "lines": [{
                "order_kind": business_kind(line),
                **{key: line.get(key, "") for key in (
                    "order_no", "child_no", "contract_no", "item_no", "customer_name",
                    "packaging_type", "paper_quality", "specification", "unit", "source_kind")},
                "quantity": line["quantity"], "received_quantity": str(line.get("received_quantity") or "0"),
                "rejected_quantity": str(line.get("rejected_quantity") or "0"), "effective_quantity": str(line.get("effective_quantity") or "0"),
            } for line in group["lines"]],
        })
    if result:
        purchase_ids = [row["id"] for row in result if row["kind"] == "PURCHASE" and not row.get("is_batch")]
        batch_ids = [row["id"] for row in result if row.get("is_batch")]
        special_version_ids = [row["id"] for row in result if row.get("special_request_id")]
        delivery_ids = [row["id"] for row in result if row["kind"] == "DELIVERY"]
        export_counts = {
            (entity_type, entity_id): count for entity_type, entity_id, count in db.execute(
                select(CartonAuditEvent.entity_type, CartonAuditEvent.entity_id,
                       func.count(CartonAuditEvent.id)).where(
                    CartonAuditEvent.factory_id == factory,
                    CartonAuditEvent.event_type == "SUPPLIER_DOCUMENT_EXPORTED",
                    or_(and_(CartonAuditEvent.entity_type == "carton_purchase_order_issue",
                            CartonAuditEvent.entity_id.in_(purchase_ids)),
                        and_(CartonAuditEvent.entity_type == "carton_special_version",
                             CartonAuditEvent.entity_id.in_(special_version_ids)),
                        and_(CartonAuditEvent.entity_type == "supplier_shipment",
                             CartonAuditEvent.entity_id.in_(delivery_ids)),
                        and_(CartonAuditEvent.entity_type == "carton_purchase_order_batch",
                             CartonAuditEvent.entity_id.in_(batch_ids))),
                ).group_by(CartonAuditEvent.entity_type, CartonAuditEvent.entity_id)).all()
        }
        for row in result:
            entity_type = _document_entity_type(row)
            row["export_count"] = export_counts.get((entity_type, row["id"]), 0)
    return sorted(result, key=lambda item: (item["created_at"], item["document_no"]), reverse=True)


def _activity_rules(db, user, factory):
    supplier = supplier_access(db, user, factory)
    order_rows = db.scalars(select(CartonOrder).where(
        CartonOrder.factory_id == factory, CartonOrder.supplier_id == supplier.id,
        CartonOrder.status.in_(VISIBLE_STATES))).all()
    order_nos = {row.id: row.order_no for row in order_rows}
    order_ids = set(order_nos)
    issue_rows = db.scalars(select(CartonPurchaseOrderIssue).where(
        CartonPurchaseOrderIssue.factory_id == factory,
        CartonPurchaseOrderIssue.order_id.in_(order_ids),
        CartonPurchaseOrderIssue.document_type != "LEGACY_BASELINE")).all()
    issue_nos = {row.id: row.document_no for row in issue_rows}
    shipments = db.scalars(select(SupplierShipment).where(
        SupplierShipment.factory_id == factory, SupplierShipment.supplier_id == supplier.id)).all()
    shipment_nos = {row.id: row.delivery_note_no for row in shipments}
    allowed = {
        "PURCHASE_ORDER_ISSUED": ("carton_purchase_order_issue", issue_nos, "采购单已生成"),
        "ORDER_DATES_CHANGED": ("carton_order", order_nos, "内部订单交期已变更"),
        "ORDER_CANCELLED": ("carton_order", order_nos, "内部订单已取消"),
        "ORDER_REDUCED": ("carton_order", order_nos, "内部订单已减单"),
        "ORDER_RETURNED": ("carton_order", order_nos, "内部订单已退单"),
        "SUPPLIER_CANCELLATION_ACKNOWLEDGED": ("carton_order", order_nos, "供应商已收到取消通知"),
        "SUPPLIER_SHIPMENT_DEFERRED": ("supplier_shipment", shipment_nos, "仓库反馈暂未到货，继续待收"),
        "SUPPLIER_SHIPMENT_NON_ARRIVAL_REOPENED": ("supplier_shipment", shipment_nos, "未到货反馈已更正，重新待收"),
        "DELIVERY_SHORTFALL_REOPENED": ("supplier_shipment", shipment_nos, "短送结束已更正，选中余量重新待收"),
        "SUPPLIER_PAPER_ACCEPTED": ("carton_order", order_nos, "纸品已确认接单"),
        "SUPPLIER_SHIPMENT_CREATED": ("supplier_shipment", shipment_nos, "供应商已确认发货"),
        "SUPPLIER_SHIPMENT_RECEIVED": ("supplier_shipment", shipment_nos, "仓库已核实送货"),
        receipt_link.EVENT: ("supplier_shipment", shipment_nos, "仓库已关联原入库凭证"),
        "SUPPLIER_SHIPMENT_NOT_RECEIVED": ("supplier_shipment", shipment_nos, "仓库反馈未收到"),
        "SUPPLIER_SHIPMENT_RECEIPT_REVERSED": ("supplier_shipment", shipment_nos, "原收料已冲销，等待仓库更正"),
        "SUPPLIER_SHIPMENT_LINE_LINKED": ("supplier_shipment", shipment_nos, "无单纸品已关联正式订单"),
    }
    return allowed


def supplier_activity_page(db, user, factory="", *, search="", event_type="", date_from="", date_to="", sort="DESC", limit=50, offset=0):
    from app.services.carton_query import date_bounds, literal_pattern
    first, after = date_bounds(date_from, date_to)
    factories = [factory] if factory else [supplier.factory_id for supplier in supplier_factories(db, user)]
    rules = {scope: _activity_rules(db, user, scope) for scope in factories}
    # Match each event to its allowed entity and supplier-owned ID before both
    # counting and paging. Never search or return the raw internal detail JSON.
    scopes = [and_(CartonAuditEvent.factory_id == scope, CartonAuditEvent.event_type == code,
                   CartonAuditEvent.entity_type == rule[0], CartonAuditEvent.entity_id.in_(rule[1]))
              for scope, allowed in rules.items() for code, rule in allowed.items()]
    query = select(CartonAuditEvent).where(or_(*scopes) if scopes else CartonAuditEvent.id.in_([]))
    if event_type:
        query = query.where(CartonAuditEvent.event_type == event_type)
    if first:
        query = query.where(CartonAuditEvent.created_at >= first)
    if after:
        query = query.where(CartonAuditEvent.created_at < after)
    if search:
        term = search.casefold()
        matches = [and_(CartonAuditEvent.factory_id == scope, CartonAuditEvent.event_type == code,
                        CartonAuditEvent.entity_id.in_([key for key, number in rule[1].items()
                            if term in number.casefold() or term in rule[2].casefold()]))
                   for scope, allowed in rules.items() for code, rule in allowed.items()]
        query = query.where(or_(CartonAuditEvent.actor_name.ilike(literal_pattern(search), escape="!"), *matches))
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    rows = db.scalars(query.order_by(CartonAuditEvent.sequence.asc() if sort == "ASC" else CartonAuditEvent.sequence.desc()).limit(limit).offset(offset))
    result = []
    for row in rows:
        rule = rules[row.factory_id][row.event_type]
        result.append({
            "id": row.id, "created_at": row.created_at, "action": rule[2],
            "reference_no": rule[1][row.entity_id], "actor_name": row.actor_name or "系统",
            "factory_id": row.factory_id, "event_type": row.event_type,
        })
    return {"items": result, "total": total, "limit": limit, "offset": offset}


def supplier_activity(db, user, factory):
    return supplier_activity_page(db, user, factory, limit=500)["items"]


def _excel_text(value):
    text = str(value or "")
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def _document_entity_type(row):
    if row.get("special_request_id"):
        return "carton_special_version"
    if row.get("is_batch"):
        return "carton_purchase_order_batch"
    return "carton_purchase_order_issue" if row["kind"] == "PURCHASE" else "supplier_shipment"


def _record_document_export(db, user, row, **details):
    _audit(db, user, row["factory_id"], "SUPPLIER_DOCUMENT_EXPORTED", _document_entity_type(row), row["id"],
           {"document_no": row["document_no"], "kind": row["kind"], **details})


def export_documents(db, user, payload: SupplierDocumentExport):
    visible_by_factory = {}
    selected = []
    for item in payload.documents:
        if item.factory_id not in visible_by_factory:
            visible_by_factory[item.factory_id] = {
                (row["kind"], row["id"]): row for row in documents(db, user, item.factory_id)
            }
        row = visible_by_factory[item.factory_id].get((item.kind, item.id))
        if row is None:
            raise HTTPException(404, "所选单据不存在或不属于当前供应商")
        selected.append(row)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for kind, title, headers in (
        ("PURCHASE", "采购单", ["厂区", "采购单号", "生成日期", "变更类型", "订单号", "客户", "合同号",
                              "客户PO", "货号", "计划交期", "纸品子单", "纸品类型", "纸质", "规格",
                              "变更前", "本次变化", "变更后", "单位", "原采购单号"]),
        ("DELIVERY", "送货单", ["送货厂区", "送货单号", "送货日期", "仓库状态", "订单号", "客户", "合同号",
                              "客户PO", "货号", "纸品子单", "纸品类型", "纸质", "规格",
                              "发货数量", "仓库实收", "单位"]),
    ):
        subset = [row for row in selected if row["kind"] == kind]
        if not subset:
            continue
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="087F74")
            cell.alignment = Alignment(wrap_text=True)
        for document in subset:
            orders = {row["order_no"]: row for row in document["orders"]}
            for line in document["lines"]:
                order = orders.get(line["order_no"], {})
                if kind == "PURCHASE":
                    values = [
                        document["factory_id"], document["document_no"], document["date"],
                        "取消采购通知：停止后续生产和送货" if document.get("is_cancellation") else "补单" if document["replenishment"] else document["document_type"],
                        line["order_no"], order.get("customer_name"), order.get("contract_no"),
                        order.get("customer_po"), order.get("item_no"), order.get("planned_date"),
                        line["child_no"], line["packaging_type"], line["paper_quality"],
                        line["specification"], line["before_quantity"], line["change_quantity"],
                        line["quantity"], line["unit"], line.get("source_document_no", document["document_no"]),
                    ]
                else:
                    values = [
                        document["factory_id"], document["document_no"], document["date"],
                        document["status"], line["order_no"], order.get("customer_name") or line.get("customer_name"),
                        order.get("contract_no") or line.get("contract_no"), order.get("customer_po"), order.get("item_no") or line.get("item_no"),
                        line["child_no"], line["packaging_type"], line["paper_quality"],
                        line["specification"], line["quantity"], line["received_quantity"], line["unit"],
                    ]
                numeric_columns = {14, 15, 16} if kind == "PURCHASE" else {13, 14}
                cells = []
                for index, value in enumerate(values):
                    if index in numeric_columns:
                        try:
                            cells.append(float(Decimal(str(value))))
                            continue
                        except (ValueError, ArithmeticError):
                            pass
                    cells.append(_excel_text(value))
                sheet.append(cells)
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(30, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    content = output.getvalue()
    for row in selected:
        _record_document_export(db, user, row)
    db.commit()
    return content


SUPPLIER_ORDER_IMPORT_HEADERS = [
    "客户单号", "客户料号", "型号", "长", "宽", "高", "数量", "材质名称", "箱型",
    "尺寸单位", "单价", "交期", "生产备注", "送货备注", "单位", "生产长", "生产宽",
    "生产高", "备注1", "备注2", "备注3",
]


def _supplier_import_dimensions(specification: str, document_no: str) -> list[float | None]:
    value = str(specification or "").strip()
    parts = re.split(r"\s*[*×xX＊]\s*", re.sub(r"\s*(?:cm|厘米|inch|in|英寸)\s*$", "", value, flags=re.IGNORECASE))
    if len(parts) not in (2, 3) or any(not re.fullmatch(r"\d+(?:\.\d+)?", part) for part in parts):
        raise HTTPException(422, f"采购单 {document_no} 的规格“{value or '空'}”无法拆成长、宽、高；请先在纸箱采购单补正规格")
    dimensions = [float(part) for part in parts]
    if any(part <= 0 for part in dimensions):
        raise HTTPException(422, f"采购单 {document_no} 的规格尺寸必须大于 0")
    return dimensions + [None] * (3 - len(dimensions))


def export_supplier_order_import(db, user, payload: SupplierDocumentExport):
    """Dongkang ERP rows use issued deltas, including B replacements, never cumulative totals."""
    if len({selection.factory_id for selection in payload.documents}) != 1:
        raise HTTPException(422, "对方导入模板一次只能包含一个送货厂区；请按厂区分别导出")
    _lock_receipt_factory(db, payload.documents[0].factory_id)
    visible_by_factory = {}
    issue_rows = []
    selected = []
    for selection in payload.documents:
        if selection.kind != "PURCHASE":
            raise HTTPException(422, "对方订单导入模板只能选择采购单")
        if selection.factory_id not in visible_by_factory:
            visible_by_factory[selection.factory_id] = {
                row["id"]: row for row in documents(db, user, selection.factory_id)
                if row["kind"] == "PURCHASE"
            }
        document = visible_by_factory[selection.factory_id].get(selection.id)
        if document is None:
            raise HTTPException(404, "所选采购单不存在或不属于当前供应商")
        if document.get("special_request_id"):
            raise HTTPException(422, "样板和杂单请下载采购单；东康正单导入模板要求真实合同及货号，不能用需求单号代填")
        if document.get("is_cancellation"):
            raise HTTPException(422, f"{document['document_no']} 是取消采购通知，请在东康 ERP 停止或关闭原单；请下载原始通知，不能作为新增订单导入")
        if document["supplier_acceptance"]["status"] != "ACCEPTED" and not payload.acknowledge_unaccepted:
            raise HTTPException(409, f"采购单 {document['document_no']}：{document['supplier_acceptance']['label']}；请先确认接单，或明确确认未接单导出提醒后再导出导入模板。导出不会自动接单")
        selected.append(document)
        if document.get("is_batch"):
            _batch, sources = load_purchase_batch(db, selection.factory_id, selection.id)
        else:
            issue = db.get(CartonPurchaseOrderIssue, selection.id)
            if issue is None or issue.factory_id != selection.factory_id:
                raise HTTPException(404, "采购单生成快照不存在")
            sources = [issue]
        issue_rows.extend((issue, _purchase_order_snapshot(issue), document) for issue in sources)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "006 (2)"
    sheet.append(SUPPLIER_ORDER_IMPORT_HEADERS)
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087F74")
    for issue, snapshot, document in issue_rows:
        header = snapshot.get("order", {})
        lines = snapshot.get("lines", [])
        if not isinstance(header, dict) or not isinstance(lines, list):
            raise HTTPException(422, f"采购单 {issue.document_no} 的生成快照不完整")
        replenishment = snapshot.get("replenishment")
        if replenishment and (not isinstance(replenishment, dict)
                              or replenishment.get("responsibility") not in {"OWN", "SUPPLIER"}):
            raise HTTPException(422, f"采购单 {issue.document_no} 的补单责任不明确，请先核对生成快照")
        changes = [Decimal(str(line.get("required_quantity_delta") or 0)) for line in lines]
        if any(change < 0 for change in changes) or not any(change > 0 for change in changes):
            raise HTTPException(422, f"采购单 {issue.document_no} 含减量或仅交期变更，不能作为新增订单导入；请单独与供应商处理变更")
        contract_no = str(header.get("contract_no") or "").strip()
        item_no = str(header.get("item_no") or "").strip()
        if not contract_no or not item_no:
            raise HTTPException(422, f"采购单 {issue.document_no} 缺少合同号或货号")
        try:
            promised_date = date.fromisoformat(str(snapshot.get("after_due_date") or ""))
        except ValueError as exc:
            raise HTTPException(422, f"采购单 {issue.document_no} 缺少有效交期") from exc
        for line, change in zip(lines, changes):
            if change <= 0:
                continue
            specification = str(line.get("specification") or "")
            declared_unit = str(line.get("dimension_unit") or "").strip().lower()
            if declared_unit and declared_unit not in {"cm", "厘米", "in", "inch", "英寸"}:
                raise HTTPException(422, f"采购单 {issue.document_no} 的尺寸单位为 {declared_unit}，请核对后再导出对方模板")
            suffix = re.search(r"(cm|厘米|inch|in|英寸)\s*$", specification, re.IGNORECASE)
            suffix_unit = ("cm" if suffix.group(1).lower() in {"cm", "厘米"} else "inch") if suffix else ""
            supplier_unit = ("cm" if declared_unit in {"cm", "厘米"} else "inch") if declared_unit else (suffix_unit or "cm")
            if suffix_unit and suffix_unit != supplier_unit:
                raise HTTPException(422, f"采购单 {issue.document_no} 的规格单位与尺寸单位不一致，请先核对")
            length, width, height = _supplier_import_dimensions(specification, issue.document_no)
            if replenishment:
                responsibility = replenishment["responsibility"]
                settlement_note = "供应商责任，免费换补" if responsibility == "SUPPLIER" else "我方责任，按实收结算"
                from app.services.carton_replenishment_supplier import commitment as b_commitment
                promise = b_commitment(db, issue.factory_id, issue.id, str(line["id"])).get("promised_date")
                promised_date = date.fromisoformat(promise or str(snapshot["after_due_date"]))
                date_note = f"本补单承诺交期 {promise}" if promise else "计划日期仅供参考，补单承诺交期待单独确认"
                delivery_note = f"补单 {issue.document_no}；{settlement_note}；送货单请注明补单号；{date_note}"
            else:
                delivery_note = ""
            unit_price = 0 if replenishment and replenishment["responsibility"] == "SUPPLIER" else float(
                Decimal(str(line.get("unit_price") or 0)))
            sheet.append([
                _excel_text(contract_no), _excel_text(item_no), _excel_text(header.get("product_name")),
                length, width, height, float(change), _excel_text(line.get("paper_quality")),
                _excel_text(line.get("packaging_type")), supplier_unit, unit_price,
                promised_date, _excel_text(line.get("note")), _excel_text(delivery_note),
                _excel_text(line.get("unit")), None, None, None, _excel_text(issue.document_no),
                _excel_text(document["document_no"]) if document.get("is_batch") else "", "",
            ])
            sheet.cell(sheet.max_row, 12).number_format = "yyyy-mm-dd"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        column[0].parent.column_dimensions[column[0].column_letter].width = min(30, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    for document in selected:
        _record_document_export(db, user, document, format="SUPPLIER_ORDER_IMPORT",
            acknowledge_unaccepted=payload.acknowledge_unaccepted,
            acceptance_status=document["supplier_acceptance"]["status"])
    db.commit()
    return output.getvalue()


def accept_line(db, user, line_id, payload: CommitmentSave):
    result = accept_lines(db, user, BatchCommitmentSave(factory_id=payload.factory_id, lines=[{
        "order_line_id": line_id, "issue_id": payload.issue_id,
        "expected_revision": payload.expected_revision, "promised_date": payload.promised_date,
    }]))
    return order_out(db, db.get(CartonOrder, result[0]))


def accept_lines(db, user, payload: BatchCommitmentSave, *, commit=True, resolve_dispute=False):
    _lock_receipt_factory(db, payload.factory_id)
    supplier = supplier_access(db, user, payload.factory_id, "carton_supplier:approve")
    validated, replacement_lines = [], []
    from app.services import carton_replenishment_supplier as replacement
    for item in payload.lines:
        line = db.get(CartonOrderLine, item.order_line_id)
        if not line or line.factory_id != payload.factory_id:
            raise HTTPException(404, "未找到纸品子单")
        order = supplier_order(db, line.order_id, payload.factory_id, supplier.id)
        issue = db.get(CartonPurchaseOrderIssue, item.issue_id)
        if issue and _purchase_order_snapshot(issue).get("replenishment"):
            replacement.validate_issue(db, order, item.issue_id, allow_dispute=resolve_dispute)
            paper = next((paper for source in replacement.sources(db, order) if source["issue_id"] == item.issue_id
                for paper in source["lines"] if paper["id"] == line.id), None)
            previous = replacement.commitment(db, payload.factory_id, item.issue_id, line.id)
            if paper is None or previous.get("revision", 0) != item.expected_revision:
                raise HTTPException(409, "补单纸品或接单版本已变化，请刷新")
            replacement_lines.append((item, order, previous))
            continue
        executable(db, order, item.issue_id, allow_dispute=resolve_dispute)
        if line.required_quantity <= 0:
            raise HTTPException(409, f"纸品 {order.order_no}/{line.line_no} 需求已取消")
        row = db.get(SupplierCommitment, line.id)
        if (row.revision if row else 0) != item.expected_revision:
            raise HTTPException(409, f"纸品 {order.order_no}/{line.line_no} 接单版本已变更，请刷新")
        validated.append((item, line, order, row))
    order_ids = []
    for item, line, order, row in validated:
        if not row:
            row = SupplierCommitment(order_line_id=line.id, factory_id=payload.factory_id, revision=0)
            db.add(row)
        previous_date = row.promised_date
        previous_issue_id = row.issue_id
        row.issue_id = item.issue_id; row.promised_date = item.promised_date.isoformat(); row.revision += 1
        row.accepted_by = user.id; row.accepted_at = now_text()
        _audit(db, user, payload.factory_id, "SUPPLIER_PAPER_ACCEPTED", "carton_order", order.id,
            {"line_id": line.id, "issue_id": row.issue_id, "promised_date": row.promised_date, "revision": row.revision, "previous_promised_date": previous_date,
             "previous_issue_id": previous_issue_id, "promise_updated": bool(previous_issue_id == row.issue_id and previous_date)})
        if order.id not in order_ids:
            order_ids.append(order.id)
    for item, order, previous in replacement_lines:
        _audit(db, user, payload.factory_id, replacement.EVENT, "carton_purchase_order_issue", item.issue_id,
            {"line_id": item.order_line_id, "issue_id": item.issue_id, "promised_date": item.promised_date.isoformat(),
             "revision": previous.get("revision", 0) + 1, "previous_promised_date": previous.get("promised_date", "")})
        if order.id not in order_ids: order_ids.append(order.id)
    db.commit() if commit else db.flush()
    return order_ids


def create_shipment(db, user, payload: ShipmentCreate, *, commit=True, source=None, allow_shared_note=False):
    _lock_receipt_factory(db, payload.factory_id)
    supplier = supplier_access(db, user, payload.factory_id, "carton_supplier:edit")
    late = payload.registration_mode == "EXISTING_RECEIPT"
    if late and (not source or payload.unmatched_lines):
        raise HTTPException(422, "后补凭证必须来自原送货单文件，且整单已匹配正式纸品")
    if payload.unmatched_lines and not source:
        raise HTTPException(422, "无单明细只能来自已预览的原始送货单")
    from app.models.carton_special_requests import SpecialShipment
    if not allow_shared_note and db.scalar(select(SpecialShipment.id).where(SpecialShipment.factory_id == payload.factory_id,
        SpecialShipment.supplier_id == supplier.id, SpecialShipment.delivery_note_no == payload.delivery_note_no)):
        raise HTTPException(409, "此送货单已关联样板或车间需求，不可再次登记正式送货")
    digest = fingerprint(payload)
    previous = db.scalar(select(SupplierShipment).where(SupplierShipment.factory_id == payload.factory_id,
        SupplierShipment.created_by == user.id, SupplierShipment.request_id == payload.request_id))
    if previous:
        if previous.supplier_id != supplier.id or previous.fingerprint != digest:
            raise HTTPException(409, "此发货提交标识已用于其他内容")
        return shipment_out(db, previous)
    from app.services.carton_supplier_evidence_guard import ensure_evidence_open
    ensure_evidence_open(db, payload.factory_id, supplier.id, payload.delivery_date)
    if db.scalar(select(SupplierShipment.id).where(SupplierShipment.factory_id == payload.factory_id,
        SupplierShipment.supplier_id == supplier.id, SupplierShipment.delivery_note_no == payload.delivery_note_no)):
        raise HTTPException(409, "此供应商送货单号已经登记")
    if not late and db.scalar(select(CartonReceipt.id).where(CartonReceipt.factory_id == payload.factory_id,
        CartonReceipt.supplier_id == supplier.id, CartonReceipt.delivery_note_no == payload.delivery_note_no)):
        raise HTTPException(409, "此送货单已有收料记录，请核对原单，不可再次登记发货")
    if payload.unmatched_lines:
        active_keys = {(_dongkang_identity(order.contract_no), _dongkang_identity(order.item_no))
            for order in db.scalars(select(CartonOrder).where(
                CartonOrder.factory_id == payload.factory_id, CartonOrder.status != "CANCELLED")).all()}
        if any(item.contract_no and (_dongkang_identity(item.contract_no),
                _dongkang_identity(item.item_no)) in active_keys for item in payload.unmatched_lines):
            raise HTTPException(409, "无单纸品已有相同合同和货号的正式订单，请重新预览并核对")
    ids = [item.order_line_id for item in payload.lines]
    authorized_ids = set(db.scalars(select(CartonOrderLine.id).join(CartonOrder,
        CartonOrder.id == CartonOrderLine.order_id).where(CartonOrderLine.id.in_(ids),
            CartonOrderLine.factory_id == payload.factory_id, CartonOrder.factory_id == payload.factory_id,
            CartonOrder.supplier_id == supplier.id, CartonOrder.status.in_(VISIBLE_STATES))).all())
    if authorized_ids != set(ids):
        raise HTTPException(404, "未找到可访问的纸品子单")
    from app.services.carton_material_correction import guard_lines
    guard_lines(db, payload.factory_id, ids)
    fulfilled = fulfilled_by_line(db, ids); reserved = outstanding(db, ids); pending = _pending_received_by_line(db, ids)
    from app.services import carton_replenishment_supplier as replacement
    from app.services.carton_replenishment_receipts import legacy_review_lines
    blocked = legacy_review_lines(db, payload.factory_id, ids)
    snapshots = []
    checked_orders = set()
    mark_manifests = {}
    for item in payload.lines:
        if item.unit_price == 0 and not item.free_confirmed:
            raise HTTPException(422, "零单价必须明确确认本行免费")
        line = db.get(CartonOrderLine, item.order_line_id)
        if not line or line.factory_id != payload.factory_id:
            raise HTTPException(404, "未找到纸品子单")
        order = supplier_order(db, line.order_id, payload.factory_id, supplier.id)
        selected_issue = db.get(CartonPurchaseOrderIssue, item.issue_id)
        is_replacement = bool(selected_issue and _purchase_order_snapshot(selected_issue).get("replenishment"))
        replacement_source = None
        if is_replacement:
            if late: raise HTTPException(409, "补单须沿本次真实送货验收，不支持关联普通历史收料")
            issue = replacement.validate_issue(db, order, item.issue_id)
            replacement_source = next(source for source in replacement.sources(db, order) if source["issue_id"] == issue.id)
            paper = next((paper for paper in replacement_source["lines"] if paper["id"] == line.id), None)
            if not paper or not all(paper["accepted"] for paper in replacement_source["lines"]):
                raise HTTPException(409, "此补单尚未整单确认接单，请先确认补单纸品及承诺交期")
            available = Decimal(paper["remaining_to_ship"])
            vendor_price = item.unit_price if item.unit_price is not None else (source or {}).get("delivery_unit_prices", {}).get(line.id)
            if replacement_source["responsibility"] == "SUPPLIER" and vendor_price is not None and Decimal(str(vendor_price)) != 0:
                raise HTTPException(422, "供应商责任补单必须免费，请核实原送货单价格")
        elif late:
            issue = latest_issue(db, order)
            if (order.status == "CANCELLED" or not issue or issue.id != item.issue_id
                or _purchase_order_pending_change(order, get_order_lines(db, order.id), issue)[0] != "NONE"):
                raise HTTPException(409, "后补凭证的采购版本已变更或订单已取消，请重新预览")
        else:
            issue = executable(db, order, item.issue_id)
        if not is_replacement and not late and order.id not in checked_orders:
            for paper in _purchase_order_snapshot(issue).get("lines", []):
                if Decimal(str(paper.get("after_required_quantity", "0"))) <= 0:
                    continue
                commitment = db.get(SupplierCommitment, str(paper["id"]))
                if not commitment or commitment.issue_id != issue.id:
                    raise HTTPException(409, f"订单 {order.order_no} 尚有纸品未确认接单，请先完成整单接单")
            checked_orders.add(order.id)
        if line.id in blocked:
            raise HTTPException(409, "此纸品有补单待核对，请联系内部仓库按原补单流程处理")
        commitment = db.get(SupplierCommitment, line.id)
        if not is_replacement and not late and (not commitment or commitment.issue_id != issue.id):
            raise HTTPException(409, "请先确认当前版本纸品子单及承诺交期")
        if not is_replacement:
            available = replacement.ordinary_capacity(db, order, line.id)
        if not late and item.quantity > available:
            raise HTTPException(409, f"纸品 {order.order_no}/{line.line_no} 超过当前未送数量")
        snapshots.append({"order_no": order.order_no, "contract_no": order.contract_no, "customer_po": order.customer_po,
            "item_no": order.item_no, "customer_name": order.customer_name, "child_no": f"{order.order_no}/{line.line_no:02d}",
            **{key: str(getattr(line, key)) for key in ("packaging_type", "paper_quality", "specification", "dimension_unit", "unit", "unit_price", "currency")},
            **({"delivery_unit_price": str(source["delivery_unit_prices"][line.id])}
                if source and line.id in source.get("delivery_unit_prices", {}) else {}),
            **({"delivery_unit_price": str(item.unit_price), "free_confirmed": item.free_confirmed}
                if item.unit_price is not None else {}),
            **({"original_source_material": source["delivery_materials"][line.id]}
                if source and line.id in source.get("delivery_materials", {}) else {}),
            **({"source_file": {key: source[key] for key in ("filename", "sha256", "rows") if key in source}}
                if source else {})})
        from app.services.carton_mark_associations import freeze_files
        if order.id not in mark_manifests: mark_manifests[order.id] = freeze_files(db, user, order)
        snapshots[-1]['file_manifest'] = mark_manifests[order.id]
        if replacement_source:
            snapshots[-1].update(replenishment_issue_id=issue.id, replenishment_document_no=issue.document_no,
                responsibility=replacement_source["responsibility"])
            if replacement_source["responsibility"] == "SUPPLIER":
                snapshots[-1].update(delivery_unit_price="0", free_confirmed=True)
    late_candidates = receipt_link.matching_receipts(db, payload.factory_id, supplier.id,
        [snapshot | {"order_line_id": item.order_line_id, "quantity": str(item.quantity)}
            for snapshot, item in zip(snapshots, payload.lines)], note_no=payload.delivery_note_no) if late else []
    if late and not late_candidates:
        raise HTTPException(409, "没有可整单关联的已入库记录，请核对原收料；后补凭证不会新建收料")
    row = SupplierShipment(id=f"CSS-{uuid4().hex}", factory_id=payload.factory_id, supplier_id=supplier.id,
        delivery_note_no=payload.delivery_note_no, delivery_date=payload.delivery_date.isoformat(), status="SENT", revision=1,
        request_id=payload.request_id, fingerprint=digest, created_by=user.id, created_at=now_text(),
        acceptance_json=json.dumps({"registration_mode": "EXISTING_RECEIPT",
            "candidate_receipt_ids": [receipt.id for receipt, _ in late_candidates]}) if late else "{}")
    db.add(row); db.flush()
    for index, item in enumerate(payload.lines):
        db.add(SupplierShipmentLine(id=f"{row.id}-{index+1:03d}", shipment_id=row.id, factory_id=payload.factory_id,
            order_line_id=item.order_line_id, issue_id=item.issue_id, quantity=item.quantity,
            snapshot_json=json.dumps(snapshots[index], ensure_ascii=False)))
    for index, item in enumerate(payload.unmatched_lines):
        snapshot = item.model_dump(mode="json") | {"order_no": "", "child_no": "", "customer_name": "",
            "customer_po": "", "source_file": source or {}}
        db.add(SupplierShipmentUnmatchedLine(id=f"{row.id}-U{index+1:03d}", shipment_id=row.id,
            factory_id=payload.factory_id, source_sheet=item.source_sheet, source_row=item.source_row,
            quantity=item.quantity, snapshot_json=json.dumps(snapshot, ensure_ascii=False)))
    _audit(db, user, payload.factory_id, "SUPPLIER_SHIPMENT_CREATED", "supplier_shipment", row.id,
        {"delivery_note_no": row.delivery_note_no, "lines": [item.model_dump(mode="json") for item in payload.lines],
         "registration_mode": payload.registration_mode,
         "unmatched_lines": [item.model_dump(mode="json") for item in payload.unmatched_lines],
         **({"source_file": source} if source else {})})
    shipment_notifications.create_notification(db, row)
    if commit:
        db.commit()
    else:
        db.flush()
    return shipment_out(db, row)


def receive_shipment(db, user, shipment_id, payload: ShipmentReceive, *, commit=True):
    internal_permission(user, payload.factory_id, "carton_procurement:receipt_write", "carton_procurement:inventory_write")
    _lock_receipt_factory(db, payload.factory_id)
    row = db.scalar(select(SupplierShipment).where(SupplierShipment.id == shipment_id, SupplierShipment.factory_id == payload.factory_id))
    if not row:
        raise HTTPException(404, "未找到此厂区的发货单")
    ensure_shipment_customers(db, user, row)
    from app.services.carton_customer_assignment import ensure_customer_operation
    for item in payload.lines:
        if item.customer_code:
            ensure_customer_operation(db, user, payload.factory_id, item.customer_code)
    digest = fingerprint(payload)
    for event in db.scalars(select(CartonAuditEvent).where(CartonAuditEvent.factory_id == payload.factory_id,
        CartonAuditEvent.entity_type == "supplier_shipment", CartonAuditEvent.entity_id == row.id,
        CartonAuditEvent.event_type.in_({"SUPPLIER_SHIPMENT_RECEIVED", "SUPPLIER_SHIPMENT_NOT_RECEIVED", "SUPPLIER_SHIPMENT_DEFERRED"}))).all():
        saved = json.loads(event.detail_json)
        if event.actor_user_id == user.id and saved.get("acceptance", {}).get("request_id") == payload.request_id:
            saved_fingerprint = saved.get("request_fingerprint") or fingerprint(ShipmentReceive.model_validate(saved["acceptance"]))
            if saved_fingerprint != digest:
                raise HTTPException(409, "此收料提交标识已用于其他内容，请核对原记录")
            return shipment_out(db, row, internal=True)
    from app.services import carton_shipment_receiving as receiving
    formal, unmatched = _shipment_sources(db, row)
    history, correction_batch, progress = receiving.progress(db, row, formal + unmatched)
    previous_receipt = correction_batch["receipt"] if correction_batch else None
    correction = previous_receipt is not None
    if receipt_link.link_only(row):
        raise HTTPException(409, "此单是后补凭证，请关联已入库记录，不能再次收料入库")
    if correction and len(payload.correction_reason.strip()) < 4:
        raise HTTPException(422, "原收料已冲销，更正验收须填写至少四字原因")
    if row.status in {"RECEIVED", "NOT_RECEIVED"} and not correction:
        if row.confirmed_by == user.id and row.confirmation_request_id == payload.request_id and row.confirmation_fingerprint == digest:
            return shipment_out(db, row, internal=True)
        raise HTTPException(409, "此发货单已确认收料，不可重复入库")
    if row.revision != payload.expected_revision:
        raise HTTPException(409, "发货单版本已变更，请刷新")
    carried = {paper.id: receiving.closure_sources(correction_batch, paper) for paper in formal + unmatched} if correction_batch else {}
    if any(carried.values()):
        from app.services.carton_supplier_evidence_guard import ensure_evidence_open
        from app.services.carton_supplier_settlement import _posting_date
        old_day = previous_receipt.acceptance_date or (_posting_date(previous_receipt.confirmed_at) if previous_receipt.confirmed_at else "")
        ensure_evidence_open(db, payload.factory_id, row.supplier_id, row.delivery_date, old_day,
            *(entry["acceptance_date"] for entries in carried.values() for entry in entries))
    if (not correction and not history and not payload.new_delivery_confirmation and any(item.received_quantity > 0 for item in payload.lines)
        and receipt_link.matching_receipts(db, payload.factory_id, row.supplier_id,
            receipt_link.shipment_papers(db, row), note_no=row.delivery_note_no, include_linked=True)):
        raise HTTPException(409, "存在相同纸品和数量的已入库记录：后补送货单请关联原记录；确为新一批到货须明确核对后确认")
    formal, unmatched = _shipment_sources(db, row)
    source = formal + unmatched
    supplied = {item.shipment_line_id: item for item in payload.lines}
    if len(supplied) != len(payload.lines) or set(supplied) != {item.id for item in source}:
        raise HTTPException(422, "必须逐条核对整张发货单，不可漏行或混入其他单据")
    active_order_keys = {(_dongkang_identity(order.contract_no), _dongkang_identity(order.item_no))
        for order in db.scalars(select(CartonOrder).where(
            CartonOrder.factory_id == payload.factory_id, CartonOrder.status != "CANCELLED")).all()} if any(
                isinstance(line, SupplierShipmentUnmatchedLine) for line in source) else set()
    reject_all = all(item.received_quantity == 0 for item in payload.lines)
    if not reject_all or payload.remaining_action == "CLOSE" or any(item.close_remaining for item in payload.lines) or correction:
        from app.services.carton_supplier_settlement import date_check, ensure_open
        date_check(payload.acceptance_date.isoformat())
        ensure_open(db, payload.factory_id, row.supplier_id, payload.acceptance_date.isoformat()[:7])
    formal_ids = [line.order_line_id for line in formal]
    fulfilled = fulfilled_by_line(db, formal_ids)
    pending = _pending_received_by_line(db, formal_ids)
    reserved = outstanding(db, formal_ids)
    own_reservations = {}
    for shipped in formal:
        own_reservations[shipped.order_line_id] = progress[shipped.id]["reserved_quantity"]
    receipt_lines = []
    receipt_sources = []
    classified_customers = {}
    original_replacements = {}
    if correction and previous_receipt:
        from app.services.carton_replenishment_receipts import receipt_links
        from app.services.carton_procurement import _receipt_lines
        links = receipt_links(db, payload.factory_id)
        for paper in _receipt_lines(db, previous_receipt.id):
            if paper.id in links and paper.order_line_id:
                original_replacements.setdefault(paper.order_line_id, set()).add(links[paper.id]['replenishment_issue_id'])
    for shipped in source:
        item = supplied[shipped.id]
        closing = payload.remaining_action == "CLOSE" or item.close_remaining
        available_quantity = progress[shipped.id]["receivable_quantity"]
        if item.received_quantity > available_quantity:
            raise HTTPException(422, "实收不能超过本送货单剩余待收数量")
        effective = item.received_quantity - item.damaged_quantity - item.rejected_quantity - item.unusable_quantity
        if effective < 0:
            raise HTTPException(422, "不可用数量不能超过实收")
        if available_quantity == 0:
            continue  # Already handled papers remain visible but cannot post again.
        if isinstance(shipped, SupplierShipmentUnmatchedLine):
            snapshot = json.loads(shipped.snapshot_json)
            if item.received_quantity == 0:
                if item.no_order_decision:
                    raise HTTPException(422, "未到货的无单纸品不能标记为样板箱或送错货")
                if closing and len(item.difference_reason.strip()) < 4:
                    raise HTTPException(422, "无单纸品未到货须填写至少四字原因")
                continue
            if item.no_order_decision == "WRONG_DELIVERY":
                if effective != 0 or len(item.difference_reason.strip()) < 4:
                    raise HTTPException(422, "送错货必须全部拒收并填写至少四字原因")
                continue
            if correction:
                previous_line = receiving.entry(correction_batch, shipped)
                if (previous_line.get("no_order_decision") not in {"SAMPLE", "MISC", "GIFT", "UNORDERED"}
                    or Decimal(str(previous_line.get("received_quantity") or 0)) <= 0):
                    raise HTTPException(422, "此无单纸品未形成原历史收料，不能借其他纸品冲销新增应付")
                if any(getattr(item, field) != previous_line.get(field, "")
                       for field in ("no_order_decision", "customer_code", "requested_by", "workshop_id")):
                    raise HTTPException(422, "历史更正须保留原纸品分类、客户、需求人及车间来源")
            elif item.no_order_decision not in {"SAMPLE", "MISC", "GIFT"}:
                raise HTTPException(422, "正单须先补建或关联正式订单；样板、杂单及免费赠送请明确选择入库类型")
            if item.workshop_id:
                from app.services.carton_master import workshop_snapshot
                workshop_snapshot(db, payload.factory_id, item.workshop_id)
            if not correction:
                prior = next((receiving.entry(batch, shipped) for batch in history
                    if batch.get("receipt") and batch["receipt"].status == "POSTED"
                    and receiving.entry(batch, shipped).get("no_order_decision") in {"SAMPLE", "MISC", "GIFT"}), None)
                if prior and any(getattr(item, key) != prior.get(key, "") for key in
                                 ("no_order_decision", "customer_code", "requested_by", "workshop_id")):
                    raise HTTPException(409, "继续收余货须保留已入库批次的类型、客户、需求人及车间")
                from app.services.carton_price_review import original_price
                supplier_price, supplier_free = original_price(shipped, "REGULAR")
                if item.no_order_decision == "GIFT":
                    if supplier_price is not None and supplier_price > 0:
                        raise HTTPException(422, "供应商原单有收费价格，请先核实赠送，不能直接改为免费")
                elif item.unit_price != (supplier_price or Decimal(0)):
                    raise HTTPException(422, "入库单价须沿用供应商送货原价，缺价请保留待供应商补价")
            if effective <= 0:
                raise HTTPException(422, "样板、杂单或赠品有效入库数量须大于 0；全部拒收请选择送错货")
            # A formal order entered after vendor import must be reconciled as a formal receipt.
            if snapshot.get("contract_no") and (
                _dongkang_identity(snapshot["contract_no"]), _dongkang_identity(snapshot["item_no"])) in active_order_keys:
                raise HTTPException(409, "此无单纸品已存在相同合同和货号的正式订单，请先核对，不得按样板箱或赠品绕过正式订单入库")
            if item.no_order_decision == "UNORDERED" and item.unit_price <= 0:
                raise HTTPException(422, "无落单入库必须填写大于零的单价")
            if item.no_order_decision == "GIFT" and item.unit_price != 0:
                raise HTTPException(422, "免费赠送的单价必须为 0")
            if not item.location_allocations:
                raise HTTPException(422, "样板、杂单或赠品有效入库须分配实际仓位")
            positions.validate_allocations(db, payload.factory_id,
                [allocation.model_dump(mode="json") for allocation in item.location_allocations], effective)
            customer_code = item.customer_code
            if customer_code.startswith("__"):
                raise HTTPException(422, "分类库存标记由原单验收生成，不可冒充客户")
            if not customer_code and item.no_order_decision in {"SAMPLE", "MISC", "GIFT"}:
                customer_code, label = {"SAMPLE": ("__SAMPLE__", "样板需求"),
                    "MISC": ("__WORKSHOP_MISC__", "车间杂用"), "GIFT": ("__SUPPLIER_GIFT__", "供应商赠送")}[item.no_order_decision]
                classified_customers[customer_code] = label
            receipt_lines.append(CartonReceiptLineCreate(source_type="AD_HOC", unordered_receipt=item.no_order_decision == "UNORDERED", dimension_unit=item.dimension_unit, customer_code=customer_code,
                contract_no=(snapshot.get("contract_no") or "") if item.no_order_decision == "UNORDERED" else (snapshot.get("contract_no") or row.delivery_note_no), item_no=snapshot["item_no"],
                packaging_type=snapshot["packaging_type"], paper_quality=snapshot["paper_quality"],
                specification=snapshot["specification"], unit=item.unit or snapshot["unit"], currency="CNY",
                delivered_quantity=available_quantity if closing else item.received_quantity, received_quantity=item.received_quantity,
                damaged_quantity=item.damaged_quantity, rejected_quantity=item.rejected_quantity,
                unusable_quantity=item.unusable_quantity, unit_price=item.unit_price,
                zero_price_kind="FREE" if item.no_order_decision == "GIFT" else item.zero_price_kind,
                location_allocations=item.location_allocations,
                feedback_note=f"{'无落单入库' if item.no_order_decision == 'UNORDERED' else ('免费赠送' if item.no_order_decision == 'GIFT' else '无正式订单杂单' if item.no_order_decision == 'MISC' else '无正式订单样板箱')}；用途：{item.sample_purpose.strip() or '未填写'}；需求人：{item.requested_by}；车间：{item.workshop_id}；{item.difference_reason}"))
            receipt_sources.append((shipped, item))
            continue
        if item.no_order_decision or item.customer_code or item.sample_purpose or item.requested_by or item.workshop_id:
            raise HTTPException(422, "正式订单纸品不得选择无单处理")
        from app.services.carton_material_correction import guard_lines
        if effective > 0:
            guard_lines(db, payload.factory_id, [shipped.order_line_id])
        line = db.get(CartonOrderLine, shipped.order_line_id)
        order = supplier_order(db, line.order_id, payload.factory_id, row.supplier_id)
        if order.status not in OPEN_STATES and item.received_quantity > 0:
            raise HTTPException(409, "关联订单当前不允许收料")
        available = line.required_quantity - fulfilled.get(line.id, 0) - pending.get(line.id, 0) - (reserved.get(line.id, 0) - own_reservations.get(line.id, 0))
        if effective > available:
            raise HTTPException(409, "有效实收超过剩余需求，其他待确认送货或收料数量已受保护")
        if effective > 0:
            if not item.location_allocations:
                raise HTTPException(422, "有效入库数量大于 0 时必须选择实际入库仓位并分配数量")
            positions.validate_allocations(db, payload.factory_id,
                [allocation.model_dump(mode="json") for allocation in item.location_allocations], effective)
        if (effective != item.received_quantity or (closing and effective != available_quantity)) and len(item.difference_reason.strip()) < 4:
            raise HTTPException(422, "短收、破损或拒收必须填写至少四字的差异原因")
        if item.received_quantity == 0:
            continue  # Non-arrival remains shipment evidence, never a fabricated receipt line.
        from app.services.carton_material_correction import latest_approved
        material_correction = latest_approved(db, payload.factory_id, line.id)
        frozen_replacement = json.loads(shipped.snapshot_json).get("replenishment_issue_id")
        if frozen_replacement and effective > 0:
            from app.services.carton_replenishment_supplier import validate_issue
            validate_issue(db, order, frozen_replacement)
        if item.replenishment_issue_id and item.replenishment_issue_id != frozen_replacement:
            raise HTTPException(409, "验收补单来源必须与原送货行一致，普通送货不能改关联补单")
        replacement_id = frozen_replacement
        if frozen_replacement and json.loads(shipped.snapshot_json).get("responsibility") == "SUPPLIER" and item.unit_price != 0:
            raise HTTPException(422, "供应商责任补单免费换补，不可填写收费价格")
        prior_ids = original_replacements.get(line.id, set())
        if prior_ids:
            if len(prior_ids) != 1 or replacement_id and replacement_id not in prior_ids:
                raise HTTPException(409, '原批次补单来源须保持一致，请核实原收料')
            replacement_id = next(iter(prior_ids))
        receipt_lines.append(CartonReceiptLineCreate(replenishment_issue_id=replacement_id, material_correction_id=material_correction['id'] if material_correction else '',
            order_line_id=line.id, delivered_quantity=available_quantity if closing else item.received_quantity,
            received_quantity=item.received_quantity, damaged_quantity=item.damaged_quantity, rejected_quantity=item.rejected_quantity,
            unusable_quantity=item.unusable_quantity, unit_price=item.unit_price,
            zero_price_kind=("FREE" if item.price_action == "FREE" else "PENDING" if item.price_action == "PENDING" and item.unit_price == 0 else item.zero_price_kind), paper_quality=item.paper_quality,
            specification=item.specification, location_allocations=item.location_allocations, feedback_note=item.difference_reason))
        receipt_sources.append((shipped, item))
    try:
        if reject_all and payload.remaining_action == "WAIT" and not any(item.close_remaining for item in payload.lines) and not correction:
            row.revision += 1
            _audit(db, user, payload.factory_id, "SUPPLIER_SHIPMENT_DEFERRED", "supplier_shipment", row.id,
                {"acceptance": payload.model_dump(mode="json"), "request_fingerprint": digest,
                 "reason": "本次暂无实到，余货继续待收", "revision": row.revision})
            if commit: db.commit()
            else: db.flush()
            return shipment_out(db, row, internal=True)
        if correction:
            row.status = "SENT"  # Also recovers reversed notes recorded before the reversal hook existed.
            row.receipt_id = previous_receipt.id
            db.flush()
        receipt = None
        if receipt_lines:
            receipt = _create_receipt(db, CartonReceiptCreate(factory_id=payload.factory_id, supplier_id=row.supplier_id,
                delivery_note_no=row.delivery_note_no, delivery_date=row.delivery_date, acceptance_date=payload.acceptance_date.isoformat(),
                note=f"供应商发货单 {row.id}，逐纸品核实实际收到，差异保留待收需求" +
                    (f"；更正原收料 {previous_receipt.receipt_no}：{payload.correction_reason}" if correction else ""),
                lines=receipt_lines), user, commit=False, supplier_shipment_id=row.id, supplier_classified_customers=classified_customers,
                corrects_receipt_id=previous_receipt.id if correction else None,
                shipment_batch_key=payload.request_id if history and not correction else None)
            confirm_receipt(db, receipt.id, CartonReceiptConfirmRequest(factory_id=payload.factory_id, expected_revision=receipt.revision, split_confirmation=payload.split_confirmation), user, commit=False)
            from app.services.carton_price_review import record_line
            from app.services.carton_procurement import _receipt_lines
            for paper, (source_paper, acceptance) in zip(_receipt_lines(db, receipt.id), receipt_sources):
                record_line(db, user, paper, source_paper, row, acceptance, source_kind="REGULAR")
        # The core lock helper expires ORM state, so update the shipment only after core posting.
        row.revision += 1
        row.confirmed_by = user.id; row.confirmed_at = now_text()
        row.confirmation_request_id = payload.request_id; row.confirmation_fingerprint = digest
        row.acceptance_json = json.dumps(payload.model_dump(mode="json"), ensure_ascii=False)
        _audit(db, user, payload.factory_id, "SUPPLIER_SHIPMENT_NOT_RECEIVED" if reject_all else "SUPPLIER_SHIPMENT_RECEIVED", "supplier_shipment", row.id,
            {"receipt_id": receipt.id if receipt else None, "acceptance": payload.model_dump(mode="json"),
             "delivery_quantities": {paper.id: str((progress[paper.id]["receivable_quantity"] if payload.remaining_action == "CLOSE" or supplied[paper.id].close_remaining
                 else supplied[paper.id].received_quantity) + sum((Decimal(entry["quantity"]) for entry in carried.get(paper.id, [])), Decimal(0))) for paper in source},
             **({"carried_closures": carried} if any(carried.values()) else {}),
             "request_fingerprint": digest, "supersedes_receipt_id": previous_receipt.id if correction else None})
        db.flush()
        batches, next_correction, updated = receiving.progress(db, row, source)
        has_pending = next_correction or any(item["remaining_quantity"] > 0 for item in updated.values())
        row.status = "SENT" if has_pending else ("RECEIVED" if any(item["received_quantity"] > 0 for item in updated.values()) else "NOT_RECEIVED")
        row.receipt_id = next_correction["receipt_id"] if next_correction else next((batch["receipt_id"] for batch in reversed(batches) if batch.get("receipt_id")), None)
        if has_pending:
            notice = shipment_notifications.create_notification(db, row)
            notice.status = "unread"; notice.handled_at = ""; notice.read_at = ""
            notice.message = "已登记本次实收，余货继续待收；请沿原送货单办理下一次入库。"
        else:
            shipment_notifications.handle_notification(db, row)
        if commit: db.commit()
        else: db.flush()
        return shipment_out(db, row, internal=True)
    except Exception:
        db.rollback()
        raise


def change_non_arrival(db, user, shipment_id, payload: ShipmentStatusChange, *, reopen=False):
    factory = require_carton_factory(payload.factory_id)
    from app.services.carton_special_requests import locked_actor
    user = locked_actor(db, user, factory)
    for permission in ("carton_procurement:read", "carton_procurement:receipt_write", "carton_procurement:inventory_write"):
        if not any(authorization_decision(user, permission, factory, department)[0] for department in CARTON_DEPARTMENTS):
            raise HTTPException(403, "没有当前厂区的供应商收料更正权限")
    row = db.scalar(select(SupplierShipment).where(SupplierShipment.id == shipment_id,
                                                 SupplierShipment.factory_id == payload.factory_id))
    if not row:
        raise HTTPException(404, "未找到此厂区的发货单")
    ensure_shipment_customers(db, user, row)
    event_type = "SUPPLIER_SHIPMENT_NON_ARRIVAL_REOPENED" if reopen else "SUPPLIER_SHIPMENT_DEFERRED"
    digest = fingerprint(payload)
    events = db.scalars(select(CartonAuditEvent).where(CartonAuditEvent.factory_id == payload.factory_id,
        CartonAuditEvent.entity_type == "supplier_shipment", CartonAuditEvent.entity_id == row.id,
        CartonAuditEvent.event_type.in_({"SUPPLIER_SHIPMENT_NON_ARRIVAL_REOPENED", "SUPPLIER_SHIPMENT_DEFERRED"}))).all()
    for event in events:
        saved = json.loads(event.detail_json)
        if event.actor_user_id == user.id and saved.get("request_id") == payload.request_id:
            if event.event_type != event_type or saved.get("request_fingerprint") != digest:
                raise HTTPException(409, "此操作标识已用于其他内容")
            return shipment_out(db, row, internal=True)
    if reopen:
        from app.services.carton_delivery_reopening import legacy_non_arrival
        return legacy_non_arrival(db, user, row, payload)
    if row.revision != payload.expected_revision:
        raise HTTPException(409, "发货单版本已变更，请刷新")
    if receipt_link.link_only(row):
        raise HTTPException(409, "已关联收料的单据请沿原收料更正流程办理")
    if row.status != "SENT":
        raise HTTPException(409, "当前送货单状态不允许此操作")
    row.revision += 1
    _audit(db, user, payload.factory_id, event_type, "supplier_shipment", row.id,
           {"reason": payload.reason, "request_id": payload.request_id, "request_fingerprint": digest,
            "revision": row.revision})
    db.commit()
    return shipment_out(db, row, internal=True)


def reopen_reversed_shipment(db, receipt, user, reason):
    """Reopen the original vendor evidence atomically with the receipt reversal."""
    from app.services import carton_shipment_receiving as receiving
    shipment = receiving.owner(db, receipt.factory_id, receipt.id)
    if not shipment:
        return
    shipment.status = "SENT"
    shipment.receipt_id = receipt.id
    shipment.revision += 1
    _audit(db, user, receipt.factory_id, "SUPPLIER_SHIPMENT_RECEIPT_REVERSED", "supplier_shipment", shipment.id,
        {"receipt_id": receipt.id, "delivery_note_no": shipment.delivery_note_no, "reason": reason,
         "acceptance": json.loads(shipment.acceptance_json or "{}")})
    notification = shipment_notifications.create_notification(db, shipment)
    notification.status = "unread"
    notification.title = f"供应商送货单 {shipment.delivery_note_no[:100]} 收料已冲销，待更正"
    notification.message = "请沿原送货单重新核实数量、价格及仓位；旧验收记录已保留。"
    notification.read_at = notification.handled_at = ""
    notification.created_at = now_text()


def link_shipment_line(db, user, shipment_id, line_id, payload: ShipmentLineLink):
    internal_permission(user, payload.factory_id, "carton_procurement:read", "carton_procurement:receipt_write", "carton_procurement:inventory_write")
    _lock_receipt_factory(db, payload.factory_id)
    shipment = db.scalar(select(SupplierShipment).where(SupplierShipment.id == shipment_id,
        SupplierShipment.factory_id == payload.factory_id))
    source = db.get(SupplierShipmentUnmatchedLine, line_id)
    if not shipment or not source or source.shipment_id != shipment.id or source.factory_id != payload.factory_id:
        raise HTTPException(404, "未找到本厂区的无单送货明细")
    ensure_shipment_customers(db, user, shipment)
    event_id = f"CAE-SHIP-LINK-{source.id}"
    previous = db.scalar(select(CartonAuditEvent).where(CartonAuditEvent.id == event_id))
    if previous:
        if previous.actor_user_id == user.id and json.loads(previous.detail_json).get("fingerprint") == fingerprint(payload):
            return shipment_out(db, shipment, internal=True)
        raise HTTPException(409, "此无单明细已经关联，不可重复或更换目标")
    prior_receipt = db.get(CartonReceipt, shipment.receipt_id) if shipment.receipt_id else None
    correction = bool(prior_receipt and prior_receipt.status == "REVERSED")
    from app.services import carton_shipment_receiving as receiving
    history, _, receipt_progress = receiving.progress(db, shipment, [source])
    closed_source = _closed_source_linkable(shipment, source, history, receipt_progress)
    if (shipment.status != "SENT" and not correction and not closed_source) or shipment.revision != payload.expected_revision or receipt_progress[source.id]["received_quantity"] > 0:
        raise HTTPException(409, "送货单状态或版本已变化，仅待验收、收料已冲销或有完整零到货关闭证据的无单行可关联")
    from app.services.carton_supplier_evidence_guard import ensure_evidence_open
    from app.services.carton_supplier_settlement import _posting_date
    ensure_evidence_open(db, payload.factory_id, shipment.supplier_id, shipment.delivery_date,
        *(batch["acceptance"].get("acceptance_date") or (batch["receipt"].acceptance_date
            if batch["receipt"] and batch["receipt"].acceptance_date else _posting_date(batch["confirmed_at"]))
            for batch in receiving.batches(db, shipment)))
    target = db.get(CartonOrderLine, payload.order_line_id)
    order = db.get(CartonOrder, target.order_id) if target else None
    if order is not None:
        from app.services.carton_customer_assignment import ensure_customer_operation
        ensure_customer_operation(db, user, payload.factory_id, order.customer_code)
    if not target or target.factory_id != payload.factory_id or not order or order.factory_id != payload.factory_id or order.supplier_id != shipment.supplier_id:
        raise HTTPException(404, "目标纸品不属于此厂区与供应商")
    from app.services.carton_procurement import get_active_customer
    customer = get_active_customer(db, payload.factory_id, payload.customer_code)
    if customer.customer_code != order.customer_code:
        raise HTTPException(422, "所选客户与正式订单客户不一致")
    if order.revision != payload.expected_order_revision:
        raise HTTPException(409, "正式订单版本已变化，请刷新后重新核对")
    issue = executable(db, order, latest_issue(db, order).id if latest_issue(db, order) else "")
    snapshot = json.loads(source.snapshot_json)
    target_material = {key: getattr(target, key) for key in ("packaging_type", "paper_quality", "specification", "dimension_unit")}
    if not snapshot.get("contract_no") or _dongkang_identity(snapshot["contract_no"]) != _dongkang_identity(order.contract_no) or _dongkang_identity(snapshot["item_no"]) != _dongkang_identity(order.item_no):
        raise HTTPException(422, "合同号或货号与正式订单不一致")
    if not all(snapshot.get(key) not in (None, "", "待复核", "纸箱") for key in ("packaging_type", "paper_quality", "specification")) or not dimension_identity(snapshot["specification"]) or material_conflicts(snapshot, target_material, dongkang_delivery=True):
        raise HTTPException(422, "纸品类型、纸质、规格或尺寸单位与正式订单不一致")
    if target.id in blocked_replacement_lines(db, payload.factory_id, [target.id]):
        raise HTTPException(409, "此纸品存在补单待核对，请按补单流程收货")
    formal, _ = _shipment_sources(db, shipment)
    if any(line.order_line_id == target.id for line in formal):
        raise HTTPException(409, "本送货单已有此纸品明细，不可重复关联")
    available = target.required_quantity - fulfilled_by_line(db, [target.id]).get(target.id, 0) - outstanding(db, [target.id]).get(target.id, 0) - _pending_received_by_line(db, [target.id]).get(target.id, 0)
    if source.quantity > available:
        raise HTTPException(409, "本次送货数量超过正式订单未收且未在途需求")
    linked_snapshot = {"order_no": order.order_no, "contract_no": order.contract_no, "customer_po": order.customer_po,
        "item_no": order.item_no, "customer_name": order.customer_name, "child_no": f"{order.order_no}/{target.line_no:02d}",
        **{key: str(getattr(target, key)) for key in ("packaging_type", "paper_quality", "specification", "dimension_unit", "unit", "unit_price", "currency")},
        **{key: str(getattr(target, key)) if getattr(target, key) is not None else None
            for key in ("net_weight_kg", "gross_weight_kg")},
        "delivery_unit_price": snapshot.get("unit_price"), "source_file": snapshot.get("source_file") or {},
        "original_unmatched_line_id": source.id, "original_source_material": snapshot}
    db.add(SupplierShipmentLine(id=f"{source.id}-L", shipment_id=shipment.id, factory_id=payload.factory_id,
        order_line_id=target.id, issue_id=issue.id, quantity=source.quantity,
        snapshot_json=json.dumps(linked_snapshot, ensure_ascii=False)))
    shipment.revision += 1
    event = _audit(db, user, payload.factory_id, "SUPPLIER_SHIPMENT_LINE_LINKED", "supplier_shipment", shipment.id,
        {"source_line_id": source.id, "source_snapshot": snapshot, "order_line_id": target.id, "issue_id": issue.id,
         "customer_code": customer.customer_code, "order_revision": order.revision, "reason": payload.reason,
         "fingerprint": fingerprint(payload)})
    event.id = event_id
    db.commit()
    return shipment_out(db, shipment, internal=True)


def link_sample_receipt(db, user, receipt_line_id: str, payload: SampleReceiptLink):
    """Reconcile a posted sample to a later formal line without posting stock or payable again."""
    internal_permission(user, payload.factory_id, "carton_procurement:order_adjust")
    _lock_receipt_factory(db, payload.factory_id)
    sample = db.get(CartonReceiptLine, receipt_line_id)
    if not sample or sample.factory_id != payload.factory_id or sample.source_type != "AD_HOC":
        raise HTTPException(404, "未找到此厂区的无单收料明细")
    receipt = db.get(CartonReceipt, sample.receipt_id)
    if not receipt or receipt.status != "POSTED" or sample.effective_quantity <= 0:
        raise HTTPException(409, "仅可关联仍有效且已入库的无单收料")
    from app.services.carton_shipment_receiving import owner
    shipment = owner(db, payload.factory_id, receipt.id)
    if not shipment:
        raise HTTPException(409, "此收料不属于供应商导入的送货单")
    if db.scalar(select(CartonAuditEvent.id).where(CartonAuditEvent.id == _sample_link_id(sample.id))):
        raise HTTPException(409, "此无单收料已经关联正式订单，不可再次关联")
    target = db.get(CartonOrderLine, payload.order_line_id)
    if not target or target.factory_id != payload.factory_id:
        raise HTTPException(404, "未找到此厂区的正式订单纸品")
    order = db.get(CartonOrder, target.order_id)
    if not order or order.factory_id != payload.factory_id or order.supplier_id != receipt.supplier_id:
        raise HTTPException(404, "目标订单不属于此厂区与供应商")
    if order.status not in OPEN_STATES or order.revision != payload.expected_order_revision:
        raise HTTPException(409, "目标订单状态或版本已变化，请刷新后重新核对")
    issue = latest_issue(db, order)
    if issue is None or _purchase_order_pending_change(order, get_order_lines(db, order.id), issue)[0] != "NONE":
        raise HTTPException(409, "目标正式订单的当前采购单尚未生成，请先生成再关联")
    matching = (sample.customer_code == order.customer_code
        and _dongkang_identity(sample.item_no) == _dongkang_identity(order.item_no)
        and _dongkang_identity(sample.packaging_type) == _dongkang_identity(target.packaging_type)
        and _dongkang_identity(sample.paper_quality) == _dongkang_identity(target.paper_quality)
        and sample.unit == target.unit
        and _dimension_identity(sample.specification)
        and _dimension_identity(sample.specification) == _dimension_identity(target.specification)
        and (_dimension_unit(sample.specification) or "cm")
            == (_dimension_unit(target.dimension_unit) or _dimension_unit(target.specification) or "cm"))
    if not matching:
        raise HTTPException(422, "客户、货号、纸品类型、材质、规格或单位与目标正式订单不一致")
    used = fulfilled_by_line(db, [target.id]).get(target.id, Decimal(0))
    in_transit = outstanding(db, [target.id]).get(target.id, Decimal(0))
    pending = _pending_received_by_line(db, [target.id]).get(target.id, Decimal(0))
    if sample.effective_quantity > target.required_quantity - used - in_transit - pending:
        raise HTTPException(409, "关联数量超过目标订单未收且未在途的需求，须先核对已发货记录")
    event = _audit(db, user, payload.factory_id, "SUPPLIER_SAMPLE_LINKED", "carton_order_line", target.id,
        {"sample_receipt_line_id": sample.id, "receipt_id": receipt.id,
         "delivery_note_no": receipt.delivery_note_no, "order_no": order.order_no,
         "quantity": str(sample.effective_quantity), "reason": payload.reason})
    event.id = _sample_link_id(sample.id)
    db.flush()
    _refresh_order_statuses(db, {order.id}, user)
    db.commit()
    return {"receipt_line_id": sample.id, "order_line_id": target.id,
            "order_no": order.order_no, "quantity": str(sample.effective_quantity)}


def validate_file(filename, content):
    if not content or len(content) > MAX_FILE:
        raise HTTPException(422, "附件不能为空且单文件最多 10 MB")
    if not filename or len(filename) > 255 or any(char in filename for char in ('/', '\\', '\r', '\n', '\x00')):
        raise HTTPException(422, "附件文件名无效")
    ext = PurePosixPath(filename).suffix.lower()
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(content))
            if not content.startswith(b"%PDF-") or reader.is_encrypted or not reader.pages:
                raise ValueError()
        except Exception:
            raise HTTPException(422, "请选择有效且未加密的 PDF 文件")
        return "application/pdf"
    formats = {".docx": ("word/document.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ".xlsx": ("xl/workbook.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ".pptx": ("ppt/presentation.xml", "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    if ext not in formats:
        raise HTTPException(422, "仅支持 PDF、DOCX、XLSX、PPTX，不支持宏或可执行文件")
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > 2000 or sum(item.file_size for item in infos) > 40 * 1024 * 1024:
                raise ValueError()
            if any(item.flag_bits & 1 or "vbaproject" in item.filename.lower() or ".." in PurePosixPath(item.filename).parts for item in infos):
                raise ValueError()
            ElementTree.fromstring(archive.read("[Content_Types].xml"))
            ElementTree.fromstring(archive.read(formats[ext][0]))
    except (BadZipFile, KeyError, ValueError, ElementTree.ParseError, RuntimeError):
        raise HTTPException(422, "附件内容不是有效的无宏 Office 文档")
    return formats[ext][1]


def upload_attachment(db, user, factory, order_id, filename, content):
    internal_permission(user, factory, "carton_procurement:order_write")
    media = validate_file(filename, content)
    _lock_receipt_factory(db, factory)
    order = db.scalar(select(CartonOrder).where(CartonOrder.id == order_id, CartonOrder.factory_id == factory))
    if not order or order.supplier_id != fixed_supplier(db, factory).id:
        raise HTTPException(404, "未找到此厂区的供应商订单")
    if order.status == "CANCELLED":
        raise HTTPException(409, "已取消订单不能新增附件")
    digest = hashlib.sha256(content).hexdigest()
    previous = db.scalar(select(SupplierAttachment).where(SupplierAttachment.order_id == order.id,
        SupplierAttachment.filename == filename).order_by(SupplierAttachment.version.desc()))
    if previous and previous.sha256 == digest:
        return attachment_out(previous)
    row = SupplierAttachment(id=f"CSA-{uuid4().hex}", factory_id=factory, order_id=order.id, filename=filename,
        media_type=media, version=(previous.version if previous else 0)+1, size=len(content), sha256=digest,
        content=content, created_by=user.id, created_at=now_text())
    db.add(row)
    _audit(db, user, factory, "SUPPLIER_ORDER_ATTACHMENT_ADDED", "carton_order", order.id,
        {"attachment_id": row.id, "filename": filename, "version": row.version, "sha256": digest})
    db.commit()
    return attachment_out(row)


def attachment_download(db, user, factory, attachment_id, *, internal=False):
    if internal:
        internal_permission(user, factory, "carton_procurement:read")
    else:
        supplier = supplier_access(db, user, factory)
    row = db.scalar(select(SupplierAttachment).where(SupplierAttachment.id == attachment_id, SupplierAttachment.factory_id == factory))
    if not row:
        raise HTTPException(404, "未找到附件")
    if not internal:
        supplier_order(db, row.order_id, factory, supplier.id)
    return row
