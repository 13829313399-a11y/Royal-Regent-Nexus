"""An optional, user-maintained carton column on Huakang A 360 export copies."""
from copy import copy
import re

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange


SHEETS = ("ITEM表", "接单表", "正单评审表")


def ensure_carton_columns(workbook, *, factory_id: str, customer_code: str):
    """Append headers only; never infer carton status or rewrite existing cells.

    Reuse existing simplified/traditional headers even if repeated: there is no
    automated value writer that needs to choose between user-owned columns.
    Preflight all sheets so an unsupported layout cannot be partially changed.
    """
    report = {"applied": False, "columns": {}, "added": []}
    if (factory_id, customer_code) != ("huakang-a", "360"):
        return report
    from app.services.customer_order_unified import CustomerOrderUnifiedError

    plans = []
    for name in SHEETS:
        if name not in workbook:
            raise CustomerOrderUnifiedError(f"360 纸箱栏：缺少 {name}")
        sheet = workbook[name]
        found = sorted(c.column for c in sheet._cells.values()
                       if c.row == 3 and c.data_type != "f"
                       and re.sub(r"\s+", "", str(c.value or "")) in {"纸箱", "紙箱"})
        if found:
            report["columns"][name] = [get_column_letter(c) for c in found]
            continue
        # Validation/table ranges may extend beyond instantiated cells. Never
        # put this manual field into such an existing user-owned column.
        # Old templates can have formatting-only cells far to the right. Those
        # are not data columns; putting the header after them creates a gap.
        right = max([28 if name == "ITEM表" else 19,
                     *(c.column for c in sheet._cells.values()
                       if c.value is not None or c.data_type == "f"
                       or c.comment is not None or c.hyperlink is not None),
                     *(r.max_col for r in sheet.merged_cells.ranges),
                     *(r.max_col for v in sheet.data_validations.dataValidation for r in v.sqref.ranges),
                     *(CellRange(t.ref).max_col for t in sheet.tables.values())])
        if right >= 16384:
            raise CustomerOrderUnifiedError(f"360 纸箱栏：{name} 最右列已占用，无法安全追加，请先核对排期")
        plans.append((sheet, right + 1))
        report["columns"][name] = [get_column_letter(right + 1)]

    for sheet, column in plans:
        common_end = 28 if sheet.title == "ITEM表" else 19
        source = max((c for c in sheet._cells.values() if c.row == 3 and c.column < column
                      and c.value not in (None, "")), key=lambda c: c.column,
                     default=sheet.cell(3, common_end))
        header = sheet.cell(3, column, "纸箱")
        header._style = copy(source._style)
        alignment = copy(source.alignment)
        alignment.wrap_text = True
        header.alignment = alignment
        sheet.column_dimensions[get_column_letter(column)].width = 18
        report["added"].append(sheet.title)
    report["applied"] = True
    return report
