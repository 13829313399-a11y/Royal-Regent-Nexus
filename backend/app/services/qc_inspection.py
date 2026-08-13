from __future__ import annotations

import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from sqlalchemy.orm.exc import StaleDataError

from app.core.time import business_now
from app.models.qc_inspection import (
    QcCustomerConfig,
    QcInspectionAuditEvent,
    QcInspectionIdempotencyRecord,
    QcInspectionOrder,
    QcInspectionProblem,
    QcInspectionReport,
    QcReportRenameBatch,
    QcReportRenameGroup,
    QcReportRenameSourceFile,
    QcScheduleChangeDecision,
    QcScheduleImportBatch,
    QcScheduleImportRow,
)
from app.schemas.qc_inspection import (
    QcCustomerConfigCreate,
    QcCustomerConfigUpdate,
    QcOrderCreate,
    QcOrderUpdate,
    QcProblemCreate,
    QcProblemUpdate,
    QcReportGenerateRequest,
    QcRenameExecuteRequest,
    QcRenamePreviewMeta,
    QcScheduleConfirmRequest,
)
from app.services.auth import ALLOWED_FACTORY_IDS, AuthContext
from app.services.legacy_excel_bridge import (
    LegacyExcelConversionError,
    convert_legacy_xls_to_xlsx,
    is_legacy_xls_workbook,
)
from app.services.qc_report_rename import (
    RENAME_RULE_VERSION,
    ReportRenameBatchError,
    ReportRenameGroup,
    ReportSourceFile,
    rename_report_batch,
)


QC_DEPARTMENTS = ("qc",)
QC_REPORT_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
MAX_SCHEDULE_IMPORT_BYTES = 20 * 1024 * 1024
SCHEDULE_IMPORT_SUFFIXES = {".xls", ".xlsx", ".xlsm"}
REPORT_TYPES = {
    "CUSTOMER_SUMMARY",
    "WEEKLY_STATISTICS",
    "HUAXING_CUSTOMER_WEEKLY_DETAIL",
    "HUAXING_WEEKLY_AGGREGATE",
    "GROUP_SUMMARY",
}
RESULT_PROBLEM_TRIGGERS = {"FAIL", "REJECTED", "CONDITIONAL_PASS", "CANCELLED"}


HEADER_ALIASES = {
    "customer_name": {"客户", "客户名称", "customer", "customername"},
    "sales_contract_no": {
        "合同号",
        "销售合同号",
        "contractno",
        "salescontractno",
        "contract",
    },
    "customer_item_no": {
        "货号",
        "客户货号",
        "item",
        "itemno",
        "customeritemno",
    },
    "customer_po_no": {
        "po",
        "pono",
        "p/o",
        "p/o#",
        "客户po",
        "客户pono",
        "采购订单号",
        "customerpo",
    },
    "product_name": {"产品名称", "产品", "品名", "product", "productname"},
    "quantity": {"数量", "订单数量", "qty", "quantity"},
    "packing": {"装箱", "装箱资料", "packing"},
    "carton_count": {"箱数", "总箱数", "cartoncount", "cartons"},
    "report_status": {"报告", "报告状态", "report", "reportstatus"},
    "production_department": {"生产部门", "生产部", "productiondepartment"},
    "export_country_code": {
        "出口国",
        "出口国家",
        "国家",
        "country",
        "exportcountry",
    },
    "shipment_date": {
        "走货期",
        "走货日期",
        "出货期",
        "出货日期",
        "shipmentdate",
        "shipdate",
    },
    "planned_inspection_date": {
        "验货期",
        "验货日期",
        "计划验货日期",
        "inspectiondate",
        "plannedinspectiondate",
    },
    "inspection_agency": {
        "验货机构",
        "验货公司",
        "inspectionagency",
        "agency",
    },
    "account_manager": {
        "跟客人员",
        "香港跟客",
        "跟单",
        "accountmanager",
        "owner",
    },
}


def now_text() -> str:
    return business_now().isoformat(timespec="seconds")


def require_qc_factory(factory_id: str, *, allow_group: bool = False) -> str:
    normalized = factory_id.strip()
    if allow_group and normalized == "*":
        return normalized
    if normalized not in ALLOWED_FACTORY_IDS:
        raise HTTPException(status_code=422, detail="QC 验货厂区无效")
    return normalized


def _canonical_json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def _payload_sha256(value: object) -> str:
    return hashlib.sha256(_canonical_json(value).encode("utf-8")).hexdigest()


def _json_dict(value: str) -> dict[str, object]:
    try:
        parsed = json.loads(value or "{}")
    except json.JSONDecodeError:
        return {}
    return parsed if isinstance(parsed, dict) else {}


def _json_list(value: str) -> list[object]:
    try:
        parsed = json.loads(value or "[]")
    except json.JSONDecodeError:
        return []
    return parsed if isinstance(parsed, list) else []


def _idempotent_response(
    db: Session,
    user: AuthContext,
    operation: str,
    request_id: str,
    payload_sha256: str,
) -> dict[str, object] | None:
    record = db.scalar(
        select(QcInspectionIdempotencyRecord).where(
            QcInspectionIdempotencyRecord.user_id == user.id,
            QcInspectionIdempotencyRecord.operation == operation,
            QcInspectionIdempotencyRecord.request_id == request_id,
        )
    )
    if record is None:
        return None
    if record.payload_sha256 != payload_sha256:
        raise HTTPException(
            status_code=409,
            detail="request_id 已用于不同的请求内容，请更换 request_id",
        )
    return _json_dict(record.response_json)


def _record_idempotency(
    db: Session,
    user: AuthContext,
    factory_id: str,
    operation: str,
    request_id: str,
    payload_sha256: str,
    response: dict[str, object],
) -> None:
    db.add(
        QcInspectionIdempotencyRecord(
            id=f"QCI-{uuid4().hex}",
            factory_id=factory_id,
            user_id=user.id,
            operation=operation,
            request_id=request_id,
            payload_sha256=payload_sha256,
            status="COMPLETED",
            response_json=_canonical_json(response),
            created_at=now_text(),
        )
    )


def _audit(
    db: Session,
    user: AuthContext,
    factory_id: str,
    event_type: str,
    entity_type: str,
    entity_id: str,
    request_id: str,
    *,
    before: dict[str, object] | None = None,
    after: dict[str, object] | None = None,
    reason: str = "",
) -> None:
    db.add(
        QcInspectionAuditEvent(
            id=f"QCA-{uuid4().hex}",
            factory_id=factory_id,
            event_type=event_type,
            entity_type=entity_type,
            entity_id=entity_id,
            before_json=_canonical_json(before or {}),
            after_json=_canonical_json(after or {}),
            reason=reason,
            request_id=request_id,
            actor_user_id=user.id,
            actor_name=user.display_name,
            created_at=now_text(),
        )
    )


def _commit(db: Session, *, conflict_message: str) -> None:
    try:
        db.commit()
    except (IntegrityError, StaleDataError) as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail=conflict_message) from exc


def _flush(db: Session, *, conflict_message: str) -> None:
    try:
        db.flush()
    except (IntegrityError, StaleDataError) as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail=conflict_message) from exc


def _order_snapshot(order: QcInspectionOrder) -> dict[str, object]:
    return {
        "id": order.id,
        "factory_id": order.factory_id,
        "inspection_no": order.inspection_no,
        "source_type": order.source_type,
        "source_import_row_id": order.source_import_row_id,
        "week_key": order.week_key,
        "customer_name": order.customer_name,
        "sales_contract_no": order.sales_contract_no,
        "customer_item_no": order.customer_item_no,
        "customer_po_no": order.customer_po_no,
        "product_name": order.product_name,
        "quantity": str(order.quantity),
        "packing": order.packing,
        "carton_count": order.carton_count,
        "report_status": order.report_status,
        "production_department": order.production_department,
        "export_country_code": order.export_country_code,
        "shipment_date": order.shipment_date,
        "planned_inspection_date": order.planned_inspection_date,
        "actual_inspection_date": order.actual_inspection_date,
        "inspection_result": order.inspection_result,
        "inspection_agency": order.inspection_agency,
        "account_manager": order.account_manager,
        "status": order.status,
        "manual_has_problem": bool(order.manual_has_problem),
        "note": order.note,
        "revision": order.revision,
        "created_by": order.created_by,
        "created_by_name": order.created_by_name,
        "updated_by": order.updated_by,
        "updated_by_name": order.updated_by_name,
        "created_at": order.created_at,
        "updated_at": order.updated_at,
    }


def problem_to_out(problem: QcInspectionProblem) -> dict[str, object]:
    return {
        "id": problem.id,
        "factory_id": problem.factory_id,
        "problem_no": problem.problem_no,
        "inspection_order_id": problem.inspection_order_id,
        "source_type": problem.source_type,
        "source_key": problem.source_key,
        "status": problem.status,
        "category": problem.category,
        "description": problem.description,
        "return_reason": problem.return_reason,
        "corrective_action": problem.corrective_action,
        "resolution": problem.resolution,
        "primary_responsible_person": problem.primary_responsible_person,
        "secondary_responsible_person": problem.secondary_responsible_person,
        "reported_date": problem.reported_date,
        "revision": problem.revision,
        "created_by": problem.created_by,
        "created_by_name": problem.created_by_name,
        "updated_by": problem.updated_by,
        "updated_by_name": problem.updated_by_name,
        "created_at": problem.created_at,
        "updated_at": problem.updated_at,
    }


def order_to_out(db: Session, order: QcInspectionOrder) -> dict[str, object]:
    problems = list(
        db.scalars(
            select(QcInspectionProblem)
            .where(
                QcInspectionProblem.factory_id == order.factory_id,
                QcInspectionProblem.inspection_order_id == order.id,
                QcInspectionProblem.status != "CANCELLED",
            )
            .order_by(QcInspectionProblem.created_at, QcInspectionProblem.id)
        ).all()
    )
    result = _order_snapshot(order)
    descriptions = [item.description for item in problems if item.description]
    result.update(
        {
            "problem_count": len(problems),
            "has_problem": bool(problems) or bool(order.manual_has_problem),
            "problem_summary": "；".join(descriptions[:3]),
        }
    )
    return result


def _get_order(db: Session, factory_id: str, order_id: str) -> QcInspectionOrder:
    order = db.get(QcInspectionOrder, order_id)
    if order is None or order.factory_id != factory_id:
        raise HTTPException(status_code=404, detail="验货主单不存在")
    return order


def _new_inspection_no() -> str:
    return f"QC-{business_now():%Y%m%d}-{uuid4().hex[:8].upper()}"


def _new_problem_no() -> str:
    return f"QCP-{business_now():%Y%m%d}-{uuid4().hex[:8].upper()}"


def _new_order(
    *,
    factory_id: str,
    source_type: str,
    source_import_row_id: str,
    fields: dict[str, object],
    user: AuthContext,
) -> QcInspectionOrder:
    timestamp = now_text()
    return QcInspectionOrder(
        id=f"QCO-{uuid4().hex}",
        factory_id=factory_id,
        inspection_no=_new_inspection_no(),
        source_type=source_type,
        source_import_row_id=source_import_row_id,
        week_key=str(fields["week_key"]),
        customer_name=str(fields["customer_name"]),
        sales_contract_no=str(fields["sales_contract_no"]),
        customer_item_no=str(fields["customer_item_no"]),
        customer_po_no=str(fields["customer_po_no"]),
        product_name=str(fields.get("product_name", "")),
        quantity=Decimal(str(fields["quantity"])),
        packing=str(fields.get("packing", "")),
        carton_count=str(fields.get("carton_count", "")),
        report_status=str(fields.get("report_status", "")),
        production_department=str(fields.get("production_department", "")),
        export_country_code=str(fields.get("export_country_code", "")),
        shipment_date=str(fields.get("shipment_date", "")),
        planned_inspection_date=str(fields.get("planned_inspection_date", "")),
        actual_inspection_date="",
        inspection_result="PENDING",
        inspection_agency=str(fields.get("inspection_agency", "")),
        account_manager=str(fields.get("account_manager", "")),
        status="SCHEDULED",
        manual_has_problem=False,
        note=str(fields.get("note", "")),
        revision=1,
        created_by=user.id,
        created_by_name=user.display_name,
        updated_by=user.id,
        updated_by_name=user.display_name,
        created_at=timestamp,
        updated_at=timestamp,
    )


def list_orders(
    db: Session,
    factory_id: str,
    *,
    week: str = "",
    search: str = "",
) -> list[dict[str, object]]:
    factory_id = require_qc_factory(factory_id)
    statement = select(QcInspectionOrder).where(QcInspectionOrder.factory_id == factory_id)
    if week:
        statement = statement.where(QcInspectionOrder.week_key == week)
    if search:
        term = f"%{search.strip()}%"
        statement = statement.where(
            or_(
                QcInspectionOrder.inspection_no.ilike(term),
                QcInspectionOrder.customer_name.ilike(term),
                QcInspectionOrder.sales_contract_no.ilike(term),
                QcInspectionOrder.customer_item_no.ilike(term),
                QcInspectionOrder.customer_po_no.ilike(term),
            )
        )
    orders = list(
        db.scalars(
            statement.order_by(
                QcInspectionOrder.planned_inspection_date,
                QcInspectionOrder.inspection_no,
            )
        ).all()
    )
    return [order_to_out(db, order) for order in orders]


def create_order(
    db: Session,
    payload: QcOrderCreate,
    user: AuthContext,
) -> dict[str, object]:
    factory_id = require_qc_factory(payload.factory_id)
    payload_data = payload.model_dump(mode="json")
    payload_hash = _payload_sha256(payload_data)
    replay = _idempotent_response(
        db, user, "ORDER_CREATE", payload.request_id, payload_hash
    )
    if replay is not None:
        return replay
    fields = payload.model_dump(exclude={"factory_id", "request_id"})
    order = _new_order(
        factory_id=factory_id,
        source_type="MANUAL",
        source_import_row_id="",
        fields=fields,
        user=user,
    )
    db.add(order)
    _flush(db, conflict_message="验货主单创建冲突，请刷新后重试")
    response = order_to_out(db, order)
    _audit(
        db,
        user,
        factory_id,
        "ORDER_CREATED",
        "qc_inspection_order",
        order.id,
        payload.request_id,
        after=response,
    )
    _record_idempotency(
        db,
        user,
        factory_id,
        "ORDER_CREATE",
        payload.request_id,
        payload_hash,
        response,
    )
    _commit(db, conflict_message="验货主单创建冲突，请刷新后重试")
    return response


def _ensure_auto_problem(
    db: Session,
    order: QcInspectionOrder,
    user: AuthContext,
    request_id: str,
) -> QcInspectionProblem | None:
    existing = db.scalar(
        select(QcInspectionProblem).where(
            QcInspectionProblem.inspection_order_id == order.id,
            QcInspectionProblem.factory_id == order.factory_id,
            QcInspectionProblem.source_key == "RESULT_TRIGGER",
        )
    )
    if existing is not None:
        return existing
    if order.inspection_result not in RESULT_PROBLEM_TRIGGERS and not order.manual_has_problem:
        return None
    timestamp = now_text()
    problem = QcInspectionProblem(
        id=f"QCPR-{uuid4().hex}",
        factory_id=order.factory_id,
        problem_no=_new_problem_no(),
        inspection_order_id=order.id,
        source_type="AUTO_RESULT",
        source_key="RESULT_TRIGGER",
        status="DRAFT",
        category="验货结果异常" if order.inspection_result in RESULT_PROBLEM_TRIGGERS else "待补充",
        description="",
        return_reason="",
        corrective_action="",
        resolution="",
        primary_responsible_person="",
        secondary_responsible_person="",
        reported_date=order.actual_inspection_date,
        revision=1,
        created_by=user.id,
        created_by_name=user.display_name,
        updated_by=user.id,
        updated_by_name=user.display_name,
        created_at=timestamp,
        updated_at=timestamp,
    )
    db.add(problem)
    _audit(
        db,
        user,
        order.factory_id,
        "PROBLEM_DRAFT_AUTO_CREATED",
        "qc_inspection_problem",
        problem.id,
        request_id,
        after=problem_to_out(problem),
        reason="验货结果非 PASS 或 QC 标记本次有问题",
    )
    return problem


ORDER_RESULT_FIELDS = {
    "actual_inspection_date",
    "inspection_result",
    "inspection_agency",
    "manual_has_problem",
    "packing",
    "carton_count",
    "report_status",
    "production_department",
}


def update_order(
    db: Session,
    order_id: str,
    payload: QcOrderUpdate,
    user: AuthContext,
) -> dict[str, object]:
    factory_id = require_qc_factory(payload.factory_id)
    payload_data = payload.model_dump(mode="json", exclude_none=True)
    payload_hash = _payload_sha256(payload_data)
    operation = f"ORDER_UPDATE:{order_id}"
    replay = _idempotent_response(db, user, operation, payload.request_id, payload_hash)
    if replay is not None:
        return replay
    order = _get_order(db, factory_id, order_id)
    if order.revision != payload.expected_revision:
        raise HTTPException(status_code=409, detail="验货主单已被其他人更新，请刷新后重试")
    before = order_to_out(db, order)
    changes = payload.model_dump(
        exclude={"factory_id", "expected_revision", "request_id", "reason"},
        exclude_none=True,
    )
    effective_actual_date = str(
        changes.get("actual_inspection_date", order.actual_inspection_date)
    )
    if (
        changes.get("inspection_result") not in {None, "PENDING"}
        and not effective_actual_date
    ):
        raise HTTPException(status_code=422, detail="提交终态验货结果必须填写实际验货日期")
    for field, value in changes.items():
        setattr(order, field, value)
    if "inspection_result" in changes:
        if order.inspection_result in {"PASS", "FAIL", "REJECTED", "CONDITIONAL_PASS"}:
            order.status = "COMPLETED"
        elif order.inspection_result == "CANCELLED":
            order.status = "CANCELLED"
        elif order.inspection_result == "PENDING":
            order.status = "SCHEDULED"
    order.revision += 1
    order.updated_by = user.id
    order.updated_by_name = user.display_name
    order.updated_at = now_text()
    _flush(db, conflict_message="验货主单已被其他人更新，请刷新后重试")
    _ensure_auto_problem(db, order, user, payload.request_id)
    _flush(db, conflict_message="验货问题草稿创建冲突，请刷新后重试")
    response = order_to_out(db, order)
    _audit(
        db,
        user,
        factory_id,
        "ORDER_UPDATED",
        "qc_inspection_order",
        order.id,
        payload.request_id,
        before=before,
        after=response,
        reason=payload.reason,
    )
    _record_idempotency(
        db,
        user,
        factory_id,
        operation,
        payload.request_id,
        payload_hash,
        response,
    )
    _commit(db, conflict_message="验货主单更新冲突，请刷新后重试")
    return response


def list_problems(
    db: Session,
    factory_id: str,
    *,
    week: str = "",
    status: str = "",
) -> list[dict[str, object]]:
    factory_id = require_qc_factory(factory_id)
    statement = (
        select(QcInspectionProblem)
        .join(
            QcInspectionOrder,
            (QcInspectionOrder.id == QcInspectionProblem.inspection_order_id)
            & (QcInspectionOrder.factory_id == QcInspectionProblem.factory_id),
        )
        .where(QcInspectionProblem.factory_id == factory_id)
    )
    if week:
        statement = statement.where(QcInspectionOrder.week_key == week)
    if status:
        statement = statement.where(QcInspectionProblem.status == status)
    problems = list(
        db.scalars(statement.order_by(QcInspectionProblem.created_at.desc())).all()
    )
    return [problem_to_out(item) for item in problems]


def create_problem(
    db: Session,
    payload: QcProblemCreate,
    user: AuthContext,
) -> dict[str, object]:
    factory_id = require_qc_factory(payload.factory_id)
    payload_data = payload.model_dump(mode="json")
    payload_hash = _payload_sha256(payload_data)
    replay = _idempotent_response(
        db, user, "PROBLEM_CREATE", payload.request_id, payload_hash
    )
    if replay is not None:
        return replay
    _get_order(db, factory_id, payload.inspection_order_id)
    timestamp = now_text()
    problem = QcInspectionProblem(
        id=f"QCPR-{uuid4().hex}",
        factory_id=factory_id,
        problem_no=_new_problem_no(),
        inspection_order_id=payload.inspection_order_id,
        source_type="MANUAL",
        source_key="",
        status=payload.status,
        category=payload.category,
        description=payload.description,
        return_reason=payload.return_reason,
        corrective_action=payload.corrective_action,
        resolution=payload.resolution,
        primary_responsible_person=payload.primary_responsible_person,
        secondary_responsible_person=payload.secondary_responsible_person,
        reported_date=payload.reported_date,
        revision=1,
        created_by=user.id,
        created_by_name=user.display_name,
        updated_by=user.id,
        updated_by_name=user.display_name,
        created_at=timestamp,
        updated_at=timestamp,
    )
    db.add(problem)
    response = problem_to_out(problem)
    _audit(
        db,
        user,
        factory_id,
        "PROBLEM_CREATED",
        "qc_inspection_problem",
        problem.id,
        payload.request_id,
        after=response,
    )
    _record_idempotency(
        db,
        user,
        factory_id,
        "PROBLEM_CREATE",
        payload.request_id,
        payload_hash,
        response,
    )
    _commit(db, conflict_message="验货问题创建冲突，请刷新后重试")
    return response


def _get_problem(db: Session, factory_id: str, problem_id: str) -> QcInspectionProblem:
    problem = db.get(QcInspectionProblem, problem_id)
    if problem is None or problem.factory_id != factory_id:
        raise HTTPException(status_code=404, detail="验货问题不存在")
    return problem


def update_problem(
    db: Session,
    problem_id: str,
    payload: QcProblemUpdate,
    user: AuthContext,
) -> dict[str, object]:
    factory_id = require_qc_factory(payload.factory_id)
    payload_data = payload.model_dump(mode="json", exclude_none=True)
    payload_hash = _payload_sha256(payload_data)
    operation = f"PROBLEM_UPDATE:{problem_id}"
    replay = _idempotent_response(db, user, operation, payload.request_id, payload_hash)
    if replay is not None:
        return replay
    problem = _get_problem(db, factory_id, problem_id)
    if problem.revision != payload.expected_revision:
        raise HTTPException(status_code=409, detail="验货问题已被其他人更新，请刷新后重试")
    before = problem_to_out(problem)
    changes = payload.model_dump(
        exclude={"factory_id", "expected_revision", "request_id", "reason"},
        exclude_none=True,
    )
    for field, value in changes.items():
        setattr(problem, field, value)
    if problem.status not in {"DRAFT", "CANCELLED"} and not problem.description.strip():
        raise HTTPException(status_code=422, detail="非草稿问题必须填写问题描述")
    problem.revision += 1
    problem.updated_by = user.id
    problem.updated_by_name = user.display_name
    problem.updated_at = now_text()
    response = problem_to_out(problem)
    _audit(
        db,
        user,
        factory_id,
        "PROBLEM_UPDATED",
        "qc_inspection_problem",
        problem.id,
        payload.request_id,
        before=before,
        after=response,
        reason=payload.reason,
    )
    _record_idempotency(
        db,
        user,
        factory_id,
        operation,
        payload.request_id,
        payload_hash,
        response,
    )
    _commit(db, conflict_message="验货问题更新冲突，请刷新后重试")
    return response


def list_customer_configs(db: Session, factory_id: str) -> list[QcCustomerConfig]:
    factory_id = require_qc_factory(factory_id)
    return list(
        db.scalars(
            select(QcCustomerConfig)
            .where(QcCustomerConfig.factory_id == factory_id)
            .order_by(QcCustomerConfig.status, QcCustomerConfig.customer_name)
        ).all()
    )


def create_customer_config(
    db: Session,
    payload: QcCustomerConfigCreate,
    user: AuthContext,
) -> dict[str, object]:
    factory_id = require_qc_factory(payload.factory_id)
    payload_data = payload.model_dump(mode="json")
    payload_hash = _payload_sha256(payload_data)
    replay = _idempotent_response(
        db, user, "CUSTOMER_CREATE", payload.request_id, payload_hash
    )
    if replay is not None:
        return replay
    timestamp = now_text()
    customer = QcCustomerConfig(
        id=f"QCC-{uuid4().hex}",
        factory_id=factory_id,
        customer_code=payload.customer_code.upper(),
        customer_name=payload.customer_name,
        is_caixing=payload.is_caixing,
        default_inspection_agency=payload.default_inspection_agency,
        default_account_manager=payload.default_account_manager,
        status="ACTIVE",
        revision=1,
        created_by=user.id,
        updated_by=user.id,
        created_at=timestamp,
        updated_at=timestamp,
    )
    db.add(customer)
    response = {
        column.name: getattr(customer, column.name)
        for column in QcCustomerConfig.__table__.columns
    }
    _audit(
        db,
        user,
        factory_id,
        "CUSTOMER_CONFIG_CREATED",
        "qc_customer_config",
        customer.id,
        payload.request_id,
        after=response,
    )
    _record_idempotency(
        db,
        user,
        factory_id,
        "CUSTOMER_CREATE",
        payload.request_id,
        payload_hash,
        response,
    )
    _commit(db, conflict_message="当前厂区已存在相同客户代码")
    return response


def update_customer_config(
    db: Session,
    customer_id: str,
    payload: QcCustomerConfigUpdate,
    user: AuthContext,
) -> dict[str, object]:
    factory_id = require_qc_factory(payload.factory_id)
    payload_data = payload.model_dump(mode="json", exclude_none=True)
    payload_hash = _payload_sha256(payload_data)
    operation = f"CUSTOMER_UPDATE:{customer_id}"
    replay = _idempotent_response(db, user, operation, payload.request_id, payload_hash)
    if replay is not None:
        return replay
    customer = db.get(QcCustomerConfig, customer_id)
    if customer is None or customer.factory_id != factory_id:
        raise HTTPException(status_code=404, detail="QC 客户配置不存在")
    if customer.revision != payload.expected_revision:
        raise HTTPException(status_code=409, detail="客户配置已被其他人更新，请刷新后重试")
    before = {
        column.name: getattr(customer, column.name)
        for column in QcCustomerConfig.__table__.columns
    }
    changes = payload.model_dump(
        exclude={"factory_id", "expected_revision", "request_id"},
        exclude_none=True,
    )
    for field, value in changes.items():
        setattr(customer, field, value)
    customer.revision += 1
    customer.updated_by = user.id
    customer.updated_at = now_text()
    response = {
        column.name: getattr(customer, column.name)
        for column in QcCustomerConfig.__table__.columns
    }
    _audit(
        db,
        user,
        factory_id,
        "CUSTOMER_CONFIG_UPDATED",
        "qc_customer_config",
        customer.id,
        payload.request_id,
        before=before,
        after=response,
    )
    _record_idempotency(
        db,
        user,
        factory_id,
        operation,
        payload.request_id,
        payload_hash,
        response,
    )
    _commit(db, conflict_message="客户配置更新冲突，请刷新后重试")
    return response


def _normalize_header(value: object) -> str:
    return re.sub(r"[\s_\-—:：#（）()]+", "", str(value or "").strip()).casefold()


NORMALIZED_HEADER_ALIASES = {
    field: {_normalize_header(alias) for alias in aliases}
    for field, aliases in HEADER_ALIASES.items()
}


def _cell_text(cell: object) -> str:
    value = getattr(cell, "value", cell)
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        number_format = str(getattr(cell, "number_format", "") or "")
        if number_format and re.fullmatch(r"0+", number_format):
            return f"{value:0{len(number_format)}d}"
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            number_format = str(getattr(cell, "number_format", "") or "")
            if number_format and re.fullmatch(r"0+", number_format):
                return f"{int(value):0{len(number_format)}d}"
            return str(int(value))
        return format(value, "f").rstrip("0").rstrip(".")
    return str(value).strip()


def _date_text(cell: object, workbook_epoch: object) -> str:
    value = getattr(cell, "value", cell)
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, workbook_epoch)
        except (ValueError, TypeError, OverflowError):
            converted = None
        if isinstance(converted, datetime):
            return converted.date().isoformat()
        if isinstance(converted, date):
            return converted.isoformat()
    text_value = _cell_text(cell)
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%Y.%m.%d"):
        try:
            return datetime.strptime(text_value, fmt).date().isoformat()
        except ValueError:
            continue
    return text_value


def _find_header(sheet: object) -> tuple[int, dict[str, int]]:
    best: tuple[int, dict[str, int]] | None = None
    for row_no, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 30)), start=1):
        mapping: dict[str, int] = {}
        for column_no, cell in enumerate(row, start=1):
            normalized = _normalize_header(cell.value)
            if not normalized:
                continue
            for field, aliases in NORMALIZED_HEADER_ALIASES.items():
                if normalized in aliases and field not in mapping:
                    mapping[field] = column_no
        if best is None or len(mapping) > len(best[1]):
            best = (row_no, mapping)
    if best is None or not {
        "customer_name",
        "sales_contract_no",
        "customer_item_no",
        "customer_po_no",
        "quantity",
    } <= set(best[1]):
        raise HTTPException(
            status_code=422,
            detail="排期表未找到客户、合同号、货号、PO、数量等必需表头",
        )
    return best


def _parse_quantity(text_value: str) -> Decimal | None:
    normalized = text_value.replace(",", "").strip()
    if not normalized:
        return None
    try:
        return Decimal(normalized)
    except InvalidOperation:
        return None


def _parse_schedule_rows(content: bytes) -> list[dict[str, object]]:
    workbook_content = content
    if is_legacy_xls_workbook(content):
        try:
            workbook_content = convert_legacy_xls_to_xlsx(content)
        except LegacyExcelConversionError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
    try:
        workbook = load_workbook(BytesIO(workbook_content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="无法读取排期 Excel 文件") from exc
    try:
        best_sheet = None
        best_header = None
        for sheet in workbook.worksheets:
            try:
                header = _find_header(sheet)
            except HTTPException:
                continue
            if best_header is None or len(header[1]) > len(best_header[1]):
                best_sheet = sheet
                best_header = header
        if best_sheet is None or best_header is None:
            raise HTTPException(
                status_code=422,
                detail="所有 Sheet 均未找到完整的 QC 排期表头",
            )
        header_row, mapping = best_header
        rows: list[dict[str, object]] = []
        for source_row_no, cells in enumerate(
            best_sheet.iter_rows(min_row=header_row + 1),
            start=header_row + 1,
        ):
            if not any(_cell_text(cell) for cell in cells):
                continue
            values: dict[str, object] = {}
            for field, column_no in mapping.items():
                cell = cells[column_no - 1] if column_no <= len(cells) else None
                if field in {"shipment_date", "planned_inspection_date"}:
                    values[field] = _date_text(cell, workbook.epoch) if cell is not None else ""
                elif field == "quantity":
                    values[field] = _parse_quantity(_cell_text(cell)) if cell is not None else None
                else:
                    values[field] = _cell_text(cell) if cell is not None else ""
            values.setdefault("product_name", "")
            values.setdefault("packing", "")
            values.setdefault("carton_count", "")
            values.setdefault("report_status", "")
            values.setdefault("production_department", "")
            values.setdefault("export_country_code", "")
            values.setdefault("shipment_date", "")
            values.setdefault("planned_inspection_date", "")
            values.setdefault("inspection_agency", "")
            values.setdefault("account_manager", "")
            values["export_country_code"] = str(values["export_country_code"]).upper()
            values["source_row_no"] = source_row_no
            rows.append(values)
        if not rows:
            raise HTTPException(status_code=422, detail="排期表没有可导入的数据行")
        if len(rows) > 5000:
            raise HTTPException(status_code=422, detail="单次排期导入不能超过 5000 行")
        return rows
    finally:
        workbook.close()


def _schedule_row_to_out(row: QcScheduleImportRow) -> dict[str, object]:
    raw = _json_dict(row.raw_json)
    return {
        "id": row.id,
        "factory_id": row.factory_id,
        "batch_id": row.batch_id,
        "source_row_no": row.source_row_no,
        "customer_name": row.customer_name,
        "sales_contract_no": row.sales_contract_no,
        "customer_item_no": row.customer_item_no,
        "customer_po_no": row.customer_po_no,
        "product_name": row.product_name,
        "quantity": str(row.quantity) if row.quantity is not None else None,
        "packing": row.packing,
        "carton_count": row.carton_count,
        "report_status": row.report_status,
        "production_department": row.production_department,
        "export_country_code": row.export_country_code,
        "shipment_date": row.shipment_date,
        "planned_inspection_date": row.planned_inspection_date,
        "inspection_agency": row.inspection_agency,
        "account_manager": row.account_manager,
        "match_status": row.match_status,
        "candidate_order_ids": [str(item) for item in _json_list(row.candidate_order_ids_json)],
        "candidate_order_revisions": {
            str(key): int(value)
            for key, value in _json_dict(row.candidate_order_revisions_json).items()
        },
        "candidate_orders": [
            item for item in raw.get("candidate_orders", []) if isinstance(item, dict)
        ],
        "changes": [item for item in _json_list(row.changes_json) if isinstance(item, dict)],
        "validation_errors": [str(item) for item in _json_list(row.validation_errors_json)],
        "decision_status": row.decision_status,
        "linked_order_id": row.linked_order_id,
    }


def batch_to_out(db: Session, batch: QcScheduleImportBatch) -> dict[str, object]:
    rows = list(
        db.scalars(
            select(QcScheduleImportRow)
            .where(
                QcScheduleImportRow.factory_id == batch.factory_id,
                QcScheduleImportRow.batch_id == batch.id,
            )
            .order_by(QcScheduleImportRow.source_row_no)
        ).all()
    )
    return {
        "id": batch.id,
        "factory_id": batch.factory_id,
        "week_key": batch.week_key,
        "source_file_name": batch.source_file_name,
        "source_file_sha256": batch.source_file_sha256,
        "source_size_bytes": batch.source_size_bytes,
        "parser_version": batch.parser_version,
        "status": batch.status,
        "row_count": batch.row_count,
        "blocking_count": batch.blocking_count,
        "summary": _json_dict(batch.summary_json),
        "revision": batch.revision,
        "created_at": batch.created_at,
        "updated_at": batch.updated_at,
        "rows": [_schedule_row_to_out(row) for row in rows],
    }


def preview_schedule_import(
    db: Session,
    *,
    factory_id: str,
    week_key: str,
    request_id: str,
    file_name: str,
    content: bytes,
    user: AuthContext,
) -> dict[str, object]:
    factory_id = require_qc_factory(factory_id)
    suffix = Path(file_name).suffix.lower()
    if suffix not in SCHEDULE_IMPORT_SUFFIXES:
        raise HTTPException(status_code=422, detail="排期导入仅支持 .xls/.xlsx/.xlsm")
    if not content:
        raise HTTPException(status_code=422, detail="排期文件不能为空")
    if len(content) > MAX_SCHEDULE_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="排期文件不能超过 20 MB")
    source_hash = hashlib.sha256(content).hexdigest()
    payload_hash = _payload_sha256(
        {
            "factory_id": factory_id,
            "week_key": week_key,
            "file_name": file_name,
            "source_sha256": source_hash,
        }
    )
    replay = _idempotent_response(
        db, user, "SCHEDULE_PREVIEW", request_id, payload_hash
    )
    if replay is not None:
        return replay
    parsed_rows = _parse_schedule_rows(content)
    timestamp = now_text()
    batch = QcScheduleImportBatch(
        id=f"QSB-{uuid4().hex}",
        factory_id=factory_id,
        week_key=week_key,
        source_file_name=Path(file_name).name,
        source_file_sha256=source_hash,
        source_size_bytes=len(content),
        parser_version="qc-schedule-v1",
        status="PREVIEW",
        row_count=len(parsed_rows),
        blocking_count=0,
        summary_json="{}",
        revision=1,
        preview_user_id=user.id,
        preview_request_id=request_id,
        preview_payload_sha256=payload_hash,
        created_by_name=user.display_name,
        created_at=timestamp,
        updated_at=timestamp,
    )
    db.add(batch)
    counts: Counter[str] = Counter()
    for parsed in parsed_rows:
        validation_errors: list[str] = []
        for field, label in (
            ("customer_name", "客户"),
            ("sales_contract_no", "合同号"),
            ("customer_item_no", "货号"),
            ("customer_po_no", "PO"),
        ):
            if not str(parsed.get(field, "")).strip():
                validation_errors.append(f"{label}不能为空")
        quantity = parsed.get("quantity")
        if quantity is None or Decimal(str(quantity)) <= 0:
            validation_errors.append("数量必须大于 0")
        for field, label in (
            ("shipment_date", "走货日期"),
            ("planned_inspection_date", "计划验货日期"),
        ):
            value = str(parsed.get(field, ""))
            if value:
                try:
                    date.fromisoformat(value)
                except ValueError:
                    validation_errors.append(f"{label}格式必须为 YYYY-MM-DD")
        candidates: list[QcInspectionOrder] = []
        if not validation_errors:
            candidates = list(
                db.scalars(
                    select(QcInspectionOrder).where(
                        QcInspectionOrder.factory_id == factory_id,
                        QcInspectionOrder.sales_contract_no
                        == str(parsed["sales_contract_no"]),
                        QcInspectionOrder.customer_item_no
                        == str(parsed["customer_item_no"]),
                        QcInspectionOrder.status == "SCHEDULED",
                        QcInspectionOrder.inspection_result == "PENDING",
                    )
                ).all()
            )
        changes: list[dict[str, object]] = []
        decision_status = "PENDING"
        if validation_errors:
            match_status = "INVALID"
        elif not candidates:
            match_status = "NEW"
        elif len(candidates) > 1:
            match_status = "MULTIPLE_MATCHES"
        else:
            match_status = "EXACT"
            candidate = candidates[0]
            for field in (
                "customer_name",
                "customer_po_no",
                "product_name",
                "quantity",
                "packing",
                "carton_count",
                "report_status",
                "production_department",
                "export_country_code",
                "shipment_date",
                "planned_inspection_date",
                "inspection_agency",
                "account_manager",
            ):
                old_value = getattr(candidate, field)
                new_value = parsed.get(field, "")
                values_equal = (
                    Decimal(str(old_value)) == Decimal(str(new_value))
                    if field == "quantity" and new_value is not None
                    else str(old_value) == str(new_value)
                )
                if not values_equal:
                    changes.append(
                        {
                            "field": field,
                            "old_value": str(old_value),
                            "new_value": str(new_value),
                        }
                    )
            if not changes:
                decision_status = "UNCHANGED"
        counts[match_status] += 1
        row = QcScheduleImportRow(
            id=f"QSR-{uuid4().hex}",
            factory_id=factory_id,
            batch_id=batch.id,
            source_row_no=int(parsed["source_row_no"]),
            customer_name=str(parsed.get("customer_name", "")),
            sales_contract_no=str(parsed.get("sales_contract_no", "")),
            customer_item_no=str(parsed.get("customer_item_no", "")),
            customer_po_no=str(parsed.get("customer_po_no", "")),
            product_name=str(parsed.get("product_name", "")),
            quantity=quantity,
            packing=str(parsed.get("packing", "")),
            carton_count=str(parsed.get("carton_count", "")),
            report_status=str(parsed.get("report_status", "")),
            production_department=str(parsed.get("production_department", "")),
            export_country_code=str(parsed.get("export_country_code", "")),
            shipment_date=str(parsed.get("shipment_date", "")),
            planned_inspection_date=str(parsed.get("planned_inspection_date", "")),
            inspection_agency=str(parsed.get("inspection_agency", "")),
            account_manager=str(parsed.get("account_manager", "")),
            match_status=match_status,
            candidate_order_ids_json=_canonical_json([item.id for item in candidates]),
            candidate_order_revisions_json=_canonical_json(
                {item.id: item.revision for item in candidates}
            ),
            changes_json=_canonical_json(changes),
            validation_errors_json=_canonical_json(validation_errors),
            decision_status=decision_status,
            linked_order_id=candidates[0].id if len(candidates) == 1 else "",
            raw_json=_canonical_json(
                {
                    **parsed,
                    "candidate_orders": [
                        {
                            "id": item.id,
                            "inspection_no": item.inspection_no,
                            "customer_po_no": item.customer_po_no,
                            "week_key": item.week_key,
                            "planned_inspection_date": item.planned_inspection_date,
                            "quantity": str(item.quantity),
                            "revision": item.revision,
                        }
                        for item in candidates
                    ],
                }
            ),
        )
        db.add(row)
    batch.blocking_count = counts["INVALID"] + counts["MULTIPLE_MATCHES"]
    if not any(row.decision_status == "PENDING" for row in db.new if isinstance(row, QcScheduleImportRow)):
        batch.status = "CONFIRMED"
    batch.summary_json = _canonical_json(
        {
            "new_count": counts["NEW"],
            "exact_count": counts["EXACT"],
            "multiple_match_count": counts["MULTIPLE_MATCHES"],
            "invalid_count": counts["INVALID"],
        }
    )
    _flush(db, conflict_message="排期预览创建冲突，请刷新后重试")
    response = batch_to_out(db, batch)
    _audit(
        db,
        user,
        factory_id,
        "SCHEDULE_PREVIEW_CREATED",
        "qc_schedule_import_batch",
        batch.id,
        request_id,
        after={
            "week_key": week_key,
            "source_file_sha256": source_hash,
            "row_count": batch.row_count,
            "blocking_count": batch.blocking_count,
        },
    )
    _record_idempotency(
        db,
        user,
        factory_id,
        "SCHEDULE_PREVIEW",
        request_id,
        payload_hash,
        response,
    )
    _commit(db, conflict_message="相同 request_id 的排期预览已存在")
    return response


def get_schedule_batch(
    db: Session, factory_id: str, batch_id: str
) -> dict[str, object]:
    factory_id = require_qc_factory(factory_id)
    batch = db.get(QcScheduleImportBatch, batch_id)
    if batch is None or batch.factory_id != factory_id:
        raise HTTPException(status_code=404, detail="排期导入批次不存在")
    return batch_to_out(db, batch)


def list_schedule_batches(
    db: Session, factory_id: str, *, week: str = "", limit: int = 20
) -> list[dict[str, object]]:
    factory_id = require_qc_factory(factory_id)
    statement = select(QcScheduleImportBatch).where(
        QcScheduleImportBatch.factory_id == factory_id
    )
    if week:
        statement = statement.where(QcScheduleImportBatch.week_key == week)
    batches = list(
        db.scalars(
            statement.order_by(QcScheduleImportBatch.created_at.desc()).limit(limit)
        ).all()
    )
    return [batch_to_out(db, item) for item in batches]


def _row_fields(row: QcScheduleImportRow, week_key: str) -> dict[str, object]:
    return {
        "week_key": week_key,
        "customer_name": row.customer_name,
        "sales_contract_no": row.sales_contract_no,
        "customer_item_no": row.customer_item_no,
        "customer_po_no": row.customer_po_no,
        "product_name": row.product_name,
        "quantity": row.quantity,
        "packing": row.packing,
        "carton_count": row.carton_count,
        "report_status": row.report_status,
        "production_department": row.production_department,
        "export_country_code": row.export_country_code,
        "shipment_date": row.shipment_date,
        "planned_inspection_date": row.planned_inspection_date,
        "inspection_agency": row.inspection_agency,
        "account_manager": row.account_manager,
        "note": "",
    }


def confirm_schedule_import(
    db: Session,
    batch_id: str,
    payload: QcScheduleConfirmRequest,
    user: AuthContext,
) -> dict[str, object]:
    factory_id = require_qc_factory(payload.factory_id)
    payload_data = payload.model_dump(mode="json")
    payload_hash = _payload_sha256(payload_data)
    operation = f"SCHEDULE_CONFIRM:{batch_id}"
    replay = _idempotent_response(db, user, operation, payload.request_id, payload_hash)
    if replay is not None:
        return replay
    batch = db.get(QcScheduleImportBatch, batch_id)
    if batch is None or batch.factory_id != factory_id:
        raise HTTPException(status_code=404, detail="排期导入批次不存在")
    if batch.revision != payload.expected_revision:
        raise HTTPException(status_code=409, detail="排期预览已变化，请刷新后重试")
    if batch.status in {"CONFIRMED", "REJECTED"}:
        raise HTTPException(status_code=409, detail="排期导入批次已完成，不能再次确认")
    rows_by_id = {
        row.id: row
        for row in db.scalars(
            select(QcScheduleImportRow).where(
                QcScheduleImportRow.factory_id == factory_id,
                QcScheduleImportRow.batch_id == batch.id,
            )
        ).all()
    }
    for decision in payload.decisions:
        row = rows_by_id.get(decision.row_id)
        if row is None:
            raise HTTPException(status_code=422, detail="确认项不属于当前排期导入批次")
        if row.decision_status != "PENDING":
            raise HTTPException(status_code=409, detail=f"第 {row.source_row_no} 行已完成确认")
        allowed_actions = {
            "NEW": {"CREATE", "SKIP"},
            "EXACT": {"UPDATE", "KEEP"},
            "MULTIPLE_MATCHES": {"UPDATE", "KEEP", "SKIP"},
            "INVALID": {"SKIP"},
        }[row.match_status]
        if decision.action not in allowed_actions:
            raise HTTPException(
                status_code=422,
                detail=f"第 {row.source_row_no} 行不能执行 {decision.action}",
            )
        target_order: QcInspectionOrder | None = None
        before: dict[str, object] = {}
        after: dict[str, object] = {}
        if decision.action == "CREATE":
            if row.validation_errors_json != "[]":
                raise HTTPException(status_code=422, detail="字段校验失败的排期行不能创建主单")
            target_order = _new_order(
                factory_id=factory_id,
                source_type="SCHEDULE_IMPORT",
                source_import_row_id=row.id,
                fields=_row_fields(row, batch.week_key),
                user=user,
            )
            db.add(target_order)
            _flush(db, conflict_message="排期确认时主单已变化，请重新预览")
            after = order_to_out(db, target_order)
        elif decision.action == "UPDATE":
            candidate_ids = [str(item) for item in _json_list(row.candidate_order_ids_json)]
            candidate_revisions = {
                str(key): int(value)
                for key, value in _json_dict(row.candidate_order_revisions_json).items()
            }
            target_id = decision.target_order_id or (
                candidate_ids[0] if len(candidate_ids) == 1 else ""
            )
            if not target_id or target_id not in candidate_ids:
                raise HTTPException(
                    status_code=422,
                    detail=f"第 {row.source_row_no} 行必须明确选择候选验货主单",
                )
            target_order = _get_order(db, factory_id, target_id)
            if (
                target_order.status != "SCHEDULED"
                or target_order.inspection_result != "PENDING"
                or target_order.revision != candidate_revisions.get(target_id)
            ):
                raise HTTPException(
                    status_code=409,
                    detail=f"第 {row.source_row_no} 行候选主单已变化，请重新预览",
                )
            before = order_to_out(db, target_order)
            fields = _row_fields(row, batch.week_key)
            for field, value in fields.items():
                if field != "note":
                    setattr(target_order, field, value)
            target_order.source_import_row_id = row.id
            target_order.revision += 1
            target_order.updated_by = user.id
            target_order.updated_by_name = user.display_name
            target_order.updated_at = now_text()
            _flush(db, conflict_message="排期确认时主单已变化，请重新预览")
            after = order_to_out(db, target_order)
        elif decision.action == "KEEP":
            candidate_ids = [str(item) for item in _json_list(row.candidate_order_ids_json)]
            candidate_revisions = {
                str(key): int(value)
                for key, value in _json_dict(row.candidate_order_revisions_json).items()
            }
            target_id = decision.target_order_id or (
                candidate_ids[0] if len(candidate_ids) == 1 else ""
            )
            if row.match_status == "MULTIPLE_MATCHES" and (
                not target_id or target_id not in candidate_ids
            ):
                raise HTTPException(
                    status_code=422,
                    detail=f"第 {row.source_row_no} 行保留时必须明确选择候选验货主单",
                )
            target_order = _get_order(db, factory_id, target_id)
            if (
                target_order.status != "SCHEDULED"
                or target_order.inspection_result != "PENDING"
                or target_order.revision != candidate_revisions.get(target_id)
            ):
                raise HTTPException(
                    status_code=409,
                    detail=f"第 {row.source_row_no} 行候选主单已变化，请重新预览",
                )
            before = order_to_out(db, target_order)
            after = before
        row.decision_status = decision.action
        row.linked_order_id = target_order.id if target_order is not None else ""
        db.add(
            QcScheduleChangeDecision(
                id=f"QCD-{uuid4().hex}",
                factory_id=factory_id,
                batch_id=batch.id,
                row_id=row.id,
                action=decision.action,
                target_order_id=row.linked_order_id,
                before_json=_canonical_json(before),
                after_json=_canonical_json(after),
                reason=decision.reason,
                actor_user_id=user.id,
                actor_name=user.display_name,
                request_id=payload.request_id,
                created_at=now_text(),
            )
        )
        _audit(
            db,
            user,
            factory_id,
            f"SCHEDULE_ROW_{decision.action}",
            "qc_schedule_import_row",
            row.id,
            payload.request_id,
            before=before,
            after=after,
            reason=decision.reason,
        )
    pending_count = sum(
        1 for row in rows_by_id.values() if row.decision_status == "PENDING"
    )
    batch.status = "CONFIRMED" if pending_count == 0 else "PARTIALLY_CONFIRMED"
    batch.revision += 1
    batch.updated_at = now_text()
    _flush(db, conflict_message="排期确认冲突，请刷新后重试")
    response = batch_to_out(db, batch)
    _audit(
        db,
        user,
        factory_id,
        "SCHEDULE_BATCH_CONFIRMED"
        if batch.status == "CONFIRMED"
        else "SCHEDULE_BATCH_PARTIALLY_CONFIRMED",
        "qc_schedule_import_batch",
        batch.id,
        payload.request_id,
        after={"status": batch.status, "revision": batch.revision},
    )
    _record_idempotency(
        db,
        user,
        factory_id,
        operation,
        payload.request_id,
        payload_hash,
        response,
    )
    _commit(db, conflict_message="排期确认冲突，请刷新后重试")
    return response


def _safe_sheet_title(value: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", value).strip() or "未命名客户"
    base = base[:31]
    candidate = base
    counter = 2
    while candidate in used:
        suffix = f"_{counter}"
        candidate = f"{base[:31-len(suffix)]}{suffix}"
        counter += 1
    used.add(candidate)
    return candidate


DETAIL_HEADERS = [
    "序号",
    "日期",
    "客户",
    "合同",
    "出口国",
    "货号",
    "产品名称",
    "数量",
    "装箱",
    "箱数",
    "结果",
    "报告",
    "生产部门",
    "验货机构",
    "验货问题",
    "第一责任人",
    "第二责任人",
    "备注",
]
RESULT_LABELS = {
    "PENDING": "待验",
    "PASS": "PASS",
    "FAIL": "FAIL",
    "REJECTED": "退货",
    "CONDITIONAL_PASS": "有条件通过",
    "CANCELLED": "已取消",
}
RETURN_RESULTS = {"FAIL", "REJECTED"}
QC_FACTORY_LABELS = {
    "huaxing": "华兴",
    "huakang-a": "华康A",
    "huakang-b": "华康B",
    "huakang-c": "华康C",
    "huakang-d": "华康D",
    "huadeng": "华登",
}
THIN_GRAY = Side(style="thin", color="B7C4C8")
TABLE_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
SUMMARY_FILL = PatternFill("solid", fgColor="E2F3F0")


def _active_problems(
    problems_by_order: dict[str, list[QcInspectionProblem]], order_id: str
) -> list[QcInspectionProblem]:
    return [
        item
        for item in problems_by_order.get(order_id, [])
        if item.status != "CANCELLED"
    ]


def _problem_text(problems: list[QcInspectionProblem]) -> str:
    return "；".join(item.description for item in problems if item.description)


def _return_reason_text(problems: list[QcInspectionProblem]) -> str:
    return "；".join(item.return_reason for item in problems if item.return_reason)


def _detail_row(
    sequence: int,
    order: QcInspectionOrder,
    problems_by_order: dict[str, list[QcInspectionProblem]],
) -> list[object]:
    problems = _active_problems(problems_by_order, order.id)
    inspection_date_text = order.actual_inspection_date
    try:
        inspection_date: object = date.fromisoformat(inspection_date_text)
    except ValueError:
        inspection_date = inspection_date_text
    return [
        sequence,
        inspection_date,
        order.customer_name,
        order.sales_contract_no,
        order.export_country_code,
        order.customer_item_no,
        order.product_name,
        float(order.quantity),
        order.packing,
        order.carton_count,
        RESULT_LABELS[order.inspection_result],
        order.report_status,
        order.production_department,
        order.inspection_agency,
        _problem_text(problems),
        "；".join(
            dict.fromkeys(
                item.primary_responsible_person
                for item in problems
                if item.primary_responsible_person
            )
        ),
        "；".join(
            dict.fromkeys(
                item.secondary_responsible_person
                for item in problems
                if item.secondary_responsible_person
            )
        ),
        order.note,
    ]


def _set_widths(sheet: object, widths: list[float]) -> None:
    for column_no, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column_no)].width = width


def _style_range(
    sheet: object,
    min_row: int,
    max_row: int,
    min_col: int,
    max_col: int,
    *,
    fill: PatternFill | None = None,
    bold: bool = False,
) -> None:
    for row in sheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            cell.border = TABLE_BORDER
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            if fill is not None:
                cell.fill = fill
            if bold:
                cell.font = Font(name="Microsoft YaHei", size=10, bold=True)
            else:
                cell.font = Font(name="Microsoft YaHei", size=9)


def _configure_print(sheet: object, max_column: int, max_row: int) -> None:
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:{get_column_letter(max_column)}{max(1, max_row)}"
    sheet.freeze_panes = "A3"
    sheet.print_title_rows = "1:2"
    sheet.sheet_view.showGridLines = False


def _write_detail_sheet(
    sheet: object,
    *,
    title: str,
    orders: list[QcInspectionOrder],
    problems_by_order: dict[str, list[QcInspectionProblem]],
    widths: list[float],
    title_height: float,
    header_height: float,
    detail_height: float,
    include_total: bool = False,
) -> int:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=18)
    sheet.cell(1, 1, title)
    sheet.append(DETAIL_HEADERS)
    for sequence, order in enumerate(orders, start=1):
        sheet.append(_detail_row(sequence, order, problems_by_order))
    data_end_row = sheet.max_row
    if include_total:
        total_row = ["总计", "", "", "", "", "", "", sum(order.quantity for order in orders)]
        total_row.extend([""] * (18 - len(total_row)))
        sheet.append(total_row)
    sheet.row_dimensions[1].height = title_height
    sheet.row_dimensions[2].height = header_height
    for row_no in range(3, sheet.max_row + 1):
        sheet.row_dimensions[row_no].height = detail_height
        sheet.cell(row_no, 2).number_format = "m月d日"
        sheet.cell(row_no, 8).number_format = "#,##0.####"
    sheet.cell(1, 1).font = Font(name="Microsoft YaHei", size=16, bold=True)
    sheet.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    _style_range(sheet, 2, 2, 1, 18, fill=HEADER_FILL, bold=True)
    if data_end_row >= 3:
        _style_range(sheet, 3, data_end_row, 1, 18)
    if include_total:
        _style_range(sheet, sheet.max_row, sheet.max_row, 1, 18, fill=SUMMARY_FILL, bold=True)
    _set_widths(sheet, widths)
    sheet.auto_filter.ref = f"A2:R{max(2, data_end_row)}"
    _configure_print(sheet, 18, sheet.max_row)
    return sheet.max_row


def _return_count(orders: list[QcInspectionOrder]) -> int:
    return sum(1 for item in orders if item.inspection_result in RETURN_RESULTS)


def _return_reasons(
    orders: list[QcInspectionOrder],
    problems_by_order: dict[str, list[QcInspectionProblem]],
) -> str:
    if not any(item.inspection_result in RETURN_RESULTS for item in orders):
        return ""
    return "；".join(
        reason
        for item in orders
        if item.inspection_result in RETURN_RESULTS
        for reason in [_return_reason_text(_active_problems(problems_by_order, item.id))]
        if reason
    )


def _write_huaxing_detail_summary(
    sheet: object,
    week_key: str,
    orders: list[QcInspectionOrder],
    problems_by_order: dict[str, list[QcInspectionProblem]],
) -> None:
    widths = [
        6.5,
        7.25,
        8.38,
        13.38,
        12.88,
        12.38,
        22.88,
        7.75,
        3.38,
        10.63,
        4.25,
        3.88,
        4.13,
        3.88,
        15.75,
        9.88,
        9.13,
        15.75,
    ]
    detail_end = _write_detail_sheet(
        sheet,
        title=f"华兴各客户每周验货明细（{week_key}）",
        orders=orders,
        problems_by_order=problems_by_order,
        widths=widths,
        title_height=25.5,
        header_height=24,
        detail_height=17.25,
    )
    summary_title_row = detail_end + 3
    sheet.merge_cells(start_row=summary_title_row, start_column=1, end_row=summary_title_row, end_column=18)
    sheet.cell(summary_title_row, 1, f"华兴厂区验货统计（{week_key}）")
    sheet.cell(summary_title_row, 1).font = Font(name="Microsoft YaHei", size=13, bold=True)
    sheet.cell(summary_title_row, 1).alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[summary_title_row].height = 24
    header_row = summary_title_row + 1
    merged_headers = (
        (1, 1, "厂区"),
        (2, 3, "客户"),
        (4, 5, "本周验货次数"),
        (6, 7, "退货次数"),
        (8, 10, "退货率"),
        (11, 14, "退货原因"),
        (15, 15, "香港跟客"),
        (16, 16, "未退货返工/AOD单数"),
        (17, 18, "备注"),
    )
    for start, end, title in merged_headers:
        if end > start:
            sheet.merge_cells(start_row=header_row, start_column=start, end_row=header_row, end_column=end)
        sheet.cell(header_row, start, title)
    by_customer: dict[str, list[QcInspectionOrder]] = defaultdict(list)
    for order in orders:
        by_customer[order.customer_name].append(order)
    current_row = header_row
    for customer, items in sorted(by_customer.items()):
        current_row += 1
        returned = _return_count(items)
        values = {
            1: "华兴",
            2: customer,
            4: len(items),
            6: returned,
            8: returned / len(items) if items else 0,
            11: _return_reasons(items, problems_by_order),
            15: "；".join(sorted({item.account_manager for item in items if item.account_manager})),
            16: "",  # 当前模型没有返工/AOD 权威字段，禁止由结果推断。
            17: "",
        }
        for column, value in values.items():
            sheet.cell(current_row, column, value)
        for start, end, _ in merged_headers:
            if end > start:
                sheet.merge_cells(start_row=current_row, start_column=start, end_row=current_row, end_column=end)
        sheet.cell(current_row, 8).number_format = "0.00%"
        sheet.row_dimensions[current_row].height = 24
    current_row += 1
    total_count = len(orders)
    total_returned = _return_count(orders)
    total_values = {
        1: "总计",
        2: "",
        4: total_count,
        6: total_returned,
        8: total_returned / total_count if total_count else 0,
        11: "",
        15: "",
        16: "",
        17: "",
    }
    for column, value in total_values.items():
        sheet.cell(current_row, column, value)
    for start, end, _ in merged_headers:
        if end > start:
            sheet.merge_cells(start_row=current_row, start_column=start, end_row=current_row, end_column=end)
    sheet.cell(current_row, 8).number_format = "0.00%"
    sheet.row_dimensions[current_row].height = 24
    _style_range(sheet, header_row, header_row, 1, 18, fill=SUMMARY_FILL, bold=True)
    if current_row - 1 > header_row:
        _style_range(sheet, header_row + 1, current_row - 1, 1, 18)
    _style_range(sheet, current_row, current_row, 1, 18, fill=SUMMARY_FILL, bold=True)
    sheet.print_area = f"A1:R{current_row}"


def _write_aggregate_sheet(
    sheet: object,
    *,
    title: str,
    headers: list[str],
    rows: list[list[object]],
    widths: list[float],
    title_height: float,
    header_height: float,
    body_height: float,
    total_height: float,
) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    sheet.cell(1, 1, title)
    sheet.append(headers)
    for values in rows:
        sheet.append(values)
    detail_end = sheet.max_row
    if headers[0] == "客户":
        total_count = sum(int(item[1]) for item in rows)
        returned = sum(int(item[2]) for item in rows)
        total_row = [
            "总计",
            total_count,
            returned,
            (total_count - returned) / total_count if total_count else 0,
            "",
            "",
            "",
            "",
        ]
        ratio_column = 4
    else:
        total_count = sum(int(item[2]) for item in rows)
        returned = sum(int(item[3]) for item in rows)
        total_row = [
            "总计",
            "",
            total_count,
            returned,
            returned / total_count if total_count else 0,
            "",
            "",
            "",
        ]
        ratio_column = 5
    sheet.append(total_row)
    total_row_no = sheet.max_row
    sheet.row_dimensions[1].height = title_height
    sheet.row_dimensions[2].height = header_height
    for row_no in range(3, detail_end + 1):
        sheet.row_dimensions[row_no].height = body_height
        sheet.cell(row_no, ratio_column).number_format = "0.00%"
    sheet.row_dimensions[total_row_no].height = total_height
    sheet.cell(total_row_no, ratio_column).number_format = "0.00%"
    sheet.cell(1, 1).font = Font(name="Microsoft YaHei", size=16, bold=True)
    sheet.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    _style_range(sheet, 2, 2, 1, 8, fill=HEADER_FILL, bold=True)
    if detail_end >= 3:
        _style_range(sheet, 3, detail_end, 1, 8)
    _style_range(sheet, total_row_no, total_row_no, 1, 8, fill=SUMMARY_FILL, bold=True)
    _set_widths(sheet, widths)
    sheet.auto_filter.ref = f"A2:H{max(2, sheet.max_row)}"
    _configure_print(sheet, 8, sheet.max_row)


def _build_report_workbook(
    report_type: str,
    week_key: str,
    orders: list[QcInspectionOrder],
    problems: list[QcInspectionProblem],
    configured_customer_names: list[str] | None = None,
) -> bytes:
    workbook = Workbook()
    default_sheet = workbook.active
    problems_by_order: dict[str, list[QcInspectionProblem]] = defaultdict(list)
    for problem in problems:
        problems_by_order[problem.inspection_order_id].append(problem)
    if report_type == "CUSTOMER_SUMMARY":
        workbook.remove(default_sheet)
        used: set[str] = set()
        by_customer: dict[str, list[QcInspectionOrder]] = defaultdict(list)
        for customer_name in configured_customer_names or []:
            by_customer[customer_name]
        for order in orders:
            by_customer[order.customer_name].append(order)
        if not by_customer:
            sheet = workbook.create_sheet("无数据")
            _write_detail_sheet(
                sheet,
                title="客户验货汇总",
                orders=[],
                problems_by_order=problems_by_order,
                widths=[3.25, 8.75, 6.88, 8.38, 12.25, 9.38, 27, 4.5, 4.75, 4.63, 3.88, 3.25, 4.13, 4.38, 13.5, 10.63, 10.5, 12.75],
                title_height=35,
                header_height=30,
                detail_height=17,
                include_total=True,
            )
        for customer, items in sorted(by_customer.items()):
            sheet = workbook.create_sheet(_safe_sheet_title(customer, used))
            _write_detail_sheet(
                sheet,
                title=f"{customer}验货汇总",
                orders=items,
                problems_by_order=problems_by_order,
                widths=[3.25, 8.75, 6.88, 8.38, 12.25, 9.38, 27, 4.5, 4.75, 4.63, 3.88, 3.25, 4.13, 4.38, 13.5, 10.63, 10.5, 12.75],
                title_height=35,
                header_height=30,
                detail_height=17,
                include_total=True,
            )
    elif report_type == "HUAXING_CUSTOMER_WEEKLY_DETAIL":
        default_sheet.title = "华兴客户每周明细"
        _write_huaxing_detail_summary(default_sheet, week_key, orders, problems_by_order)
    elif report_type == "WEEKLY_STATISTICS":
        default_sheet.title = "每周统计"
        _write_detail_sheet(
            default_sheet,
            title=f"{week_key[:4]}年每周客验货情况统计表",
            orders=orders,
            problems_by_order=problems_by_order,
            widths=[3.75, 7.5, 9.38, 13.75, 13.75, 15.25, 21, 4.63, 3.63, 5.25, 4.25, 3.38, 4.25, 4.13, 15.13, 10, 10, 10.63],
            title_height=35,
            header_height=30.75,
            detail_height=17.25,
        )
    elif report_type == "HUAXING_WEEKLY_AGGREGATE":
        default_sheet.title = "华兴每周聚合"
        by_customer = defaultdict(list)
        for order in orders:
            by_customer[order.customer_name].append(order)
        rows = []
        for customer, items in sorted(by_customer.items()):
            returned = _return_count(items)
            rows.append(
                [
                    customer,
                    len(items),
                    returned,
                    (len(items) - returned) / len(items) if items else 0,
                    _return_reasons(items, problems_by_order),
                    "；".join(
                        dict.fromkeys(
                            responsible
                            for item in items
                            for problem in _active_problems(problems_by_order, item.id)
                            for responsible in (
                                problem.primary_responsible_person,
                                problem.secondary_responsible_person,
                            )
                            if responsible
                        )
                    ),
                    "；".join(
                        detail
                        for item in items
                        for detail in [
                            "；".join(
                                problem.resolution or problem.corrective_action
                                for problem in _active_problems(problems_by_order, item.id)
                                if problem.resolution or problem.corrective_action
                            )
                        ]
                        if detail
                    ),
                    "",
                ]
            )
        _write_aggregate_sheet(
            default_sheet,
            title=f"华兴每周验货汇总（{week_key}）",
            headers=["客户", "本周验货次数", "退货次数", "合格率", "退货原因", "第一/二责任人", "处理明细", "备注"],
            rows=rows,
            widths=[14.13, 13.25, 10, 9.63, 22.25, 20.5, 15.75, 16],
            title_height=44.25,
            header_height=47,
            body_height=49,
            total_height=38.25,
        )
    elif report_type == "GROUP_SUMMARY":
        default_sheet.title = "集团汇总"
        by_factory_customer: dict[tuple[str, str], list[QcInspectionOrder]] = defaultdict(list)
        for order in orders:
            by_factory_customer[(order.factory_id, order.customer_name)].append(order)
        rows = []
        for (factory_id, customer), items in sorted(by_factory_customer.items()):
            returned = _return_count(items)
            rows.append(
                [
                    QC_FACTORY_LABELS.get(factory_id, factory_id),
                    customer,
                    len(items),
                    returned,
                    returned / len(items) if items else 0,
                    _return_reasons(items, problems_by_order),
                    "；".join(sorted({item.account_manager for item in items if item.account_manager})),
                    "",
                ]
            )
        _write_aggregate_sheet(
            default_sheet,
            title=f"华登集团验货汇总（{week_key}）",
            headers=["厂区", "客户", "本周验货次数", "退货次数", "退货率", "退货原因", "香港跟客", "备注"],
            rows=rows,
            widths=[11.63, 14.88, 16.25, 11.13, 9.75, 21.38, 11.13, 13.38],
            title_height=41,
            header_height=41,
            body_height=45,
            total_height=45,
        )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def report_to_out(report: QcInspectionReport) -> dict[str, object]:
    return {
        "id": report.id,
        "factory_id": report.factory_id,
        "week_key": report.week_key,
        "report_type": report.report_type,
        "status": report.status,
        "is_formal_snapshot": bool(report.is_formal_snapshot),
        "snapshot_revision": report.snapshot_revision,
        "source_revision_sha256": report.source_revision_sha256,
        "artifact_file_name": report.artifact_file_name,
        "artifact_media_type": report.artifact_media_type,
        "artifact_sha256": report.artifact_sha256,
        "artifact_size_bytes": report.artifact_size_bytes,
        "generation_summary": _json_dict(report.generation_summary_json),
        "error_message": report.error_message,
        "requested_by": report.requested_by,
        "requested_by_name": report.requested_by_name,
        "created_at": report.created_at,
    }


def generate_report(
    db: Session,
    payload: QcReportGenerateRequest,
    user: AuthContext,
) -> dict[str, object]:
    allow_group = payload.report_type == "GROUP_SUMMARY"
    factory_id = require_qc_factory(payload.factory_id, allow_group=allow_group)
    if allow_group and factory_id != "*":
        raise HTTPException(status_code=422, detail="集团汇总必须使用 factory_id='*'")
    if not allow_group and factory_id == "*":
        raise HTTPException(status_code=422, detail="本厂报表必须指定具体厂区")
    if payload.report_type.startswith("HUAXING_") and factory_id != "huaxing":
        raise HTTPException(status_code=422, detail="华兴报表只能使用华兴厂区数据")
    payload_data = payload.model_dump(mode="json")
    payload_hash = _payload_sha256(payload_data)
    replay = _idempotent_response(db, user, "REPORT_GENERATE", payload.request_id, payload_hash)
    if replay is not None:
        return replay
    order_statement = select(QcInspectionOrder).where(
        QcInspectionOrder.week_key == payload.week_key
    )
    if factory_id != "*":
        order_statement = order_statement.where(QcInspectionOrder.factory_id == factory_id)
    scoped_orders = list(
        db.scalars(
            order_statement.order_by(
                QcInspectionOrder.factory_id,
                QcInspectionOrder.customer_name,
                QcInspectionOrder.inspection_no,
            )
        ).all()
    )
    completed_results = {"PASS", "FAIL", "REJECTED", "CONDITIONAL_PASS"}
    orders = [
        item
        for item in scoped_orders
        if item.status == "COMPLETED"
        and item.inspection_result in completed_results
        and bool(item.actual_inspection_date)
    ]
    excluded_cancelled_count = sum(
        1
        for item in scoped_orders
        if item.status == "CANCELLED" or item.inspection_result == "CANCELLED"
    )
    excluded_pending_count = len(scoped_orders) - len(orders) - excluded_cancelled_count
    order_ids = [item.id for item in orders]
    problems = (
        list(
            db.scalars(
                select(QcInspectionProblem)
                .where(QcInspectionProblem.inspection_order_id.in_(order_ids))
                .order_by(QcInspectionProblem.problem_no)
            ).all()
        )
        if order_ids
        else []
    )
    configured_customers = (
        list(
            db.scalars(
                select(QcCustomerConfig)
                .where(
                    QcCustomerConfig.factory_id == factory_id,
                    QcCustomerConfig.status == "ACTIVE",
                )
                .order_by(QcCustomerConfig.customer_name)
            ).all()
        )
        if payload.report_type == "CUSTOMER_SUMMARY" and factory_id != "*"
        else []
    )
    configured_customer_names = [item.customer_name for item in configured_customers]
    source_revision = _payload_sha256(
        {
            "orders": [(item.id, item.revision) for item in orders],
            "problems": [(item.id, item.revision) for item in problems],
            "customer_configs": [
                (item.id, item.revision, item.customer_name, item.status)
                for item in configured_customers
            ],
        }
    )
    artifact = _build_report_workbook(
        payload.report_type,
        payload.week_key,
        orders,
        problems,
        configured_customer_names,
    )
    artifact_hash = hashlib.sha256(artifact).hexdigest()
    snapshot_revision = (
        db.scalar(
            select(func.max(QcInspectionReport.snapshot_revision)).where(
                QcInspectionReport.factory_id == factory_id,
                QcInspectionReport.week_key == payload.week_key,
                QcInspectionReport.report_type == payload.report_type,
            )
        )
        or 0
    ) + 1
    file_name = (
        f"QC_{payload.report_type}_{factory_id}_{payload.week_key}_"
        f"R{snapshot_revision}.xlsx"
    )
    report = QcInspectionReport(
        id=f"QCR-{uuid4().hex}",
        factory_id=factory_id,
        week_key=payload.week_key,
        report_type=payload.report_type,
        status="GENERATED",
        is_formal_snapshot=payload.is_formal_snapshot,
        snapshot_revision=snapshot_revision,
        source_revision_sha256=source_revision,
        artifact_file_name=file_name,
        artifact_media_type=QC_REPORT_MEDIA_TYPE,
        artifact_sha256=artifact_hash,
        artifact_size_bytes=len(artifact),
        artifact_bytes=artifact,
        generation_summary_json=_canonical_json(
            {
                "order_count": len(orders),
                "problem_count": len(problems),
                "customer_config_count": len(configured_customers),
                "excluded_pending_count": excluded_pending_count,
                "excluded_cancelled_count": excluded_cancelled_count,
                "source": "qc_inspection_orders+qc_inspection_problems",
            }
        ),
        error_message="",
        requested_by=user.id,
        requested_by_name=user.display_name,
        created_at=now_text(),
    )
    db.add(report)
    response = report_to_out(report)
    _audit(
        db,
        user,
        factory_id,
        "REPORT_GENERATED",
        "qc_inspection_report",
        report.id,
        payload.request_id,
        after=response,
    )
    _record_idempotency(
        db,
        user,
        factory_id,
        "REPORT_GENERATE",
        payload.request_id,
        payload_hash,
        response,
    )
    _commit(db, conflict_message="报表生成请求冲突，请更换 request_id")
    return response


def list_reports(
    db: Session,
    factory_id: str,
    *,
    week: str = "",
) -> list[dict[str, object]]:
    factory_id = require_qc_factory(factory_id, allow_group=True)
    statement = select(QcInspectionReport).where(
        QcInspectionReport.factory_id == factory_id
    )
    if week:
        statement = statement.where(QcInspectionReport.week_key == week)
    reports = list(
        db.scalars(statement.order_by(QcInspectionReport.created_at.desc())).all()
    )
    return [report_to_out(item) for item in reports]


def get_report_artifact(
    db: Session, factory_id: str, report_id: str
) -> tuple[str, str, bytes]:
    factory_id = require_qc_factory(factory_id, allow_group=True)
    report = db.get(QcInspectionReport, report_id)
    if report is None or report.factory_id != factory_id:
        raise HTTPException(status_code=404, detail="QC 报表不存在")
    if report.status != "GENERATED" or not report.artifact_bytes:
        raise HTTPException(status_code=409, detail="QC 报表尚无可下载工件")
    artifact = bytes(report.artifact_bytes)
    if (
        len(artifact) != report.artifact_size_bytes
        or hashlib.sha256(artifact).hexdigest() != report.artifact_sha256
    ):
        raise HTTPException(status_code=409, detail="QC 报表工件完整性校验失败")
    return report.artifact_file_name, report.artifact_media_type, artifact


def list_audit_events(
    db: Session,
    factory_id: str,
    *,
    limit: int = 100,
) -> list[dict[str, object]]:
    factory_id = require_qc_factory(factory_id)
    events = list(
        db.scalars(
            select(QcInspectionAuditEvent)
            .where(QcInspectionAuditEvent.factory_id == factory_id)
            .order_by(QcInspectionAuditEvent.created_at.desc())
            .limit(limit)
        ).all()
    )
    return [
        {
            "id": event.id,
            "factory_id": event.factory_id,
            "event_type": event.event_type,
            "entity_type": event.entity_type,
            "entity_id": event.entity_id,
            "before": _json_dict(event.before_json),
            "after": _json_dict(event.after_json),
            "reason": event.reason,
            "request_id": event.request_id,
            "actor_user_id": event.actor_user_id,
            "actor_name": event.actor_name,
            "created_at": event.created_at,
        }
        for event in events
    ]


def workspace(
    db: Session,
    factory_id: str,
    week: str,
) -> dict[str, object]:
    factory_id = require_qc_factory(factory_id)
    orders = list_orders(db, factory_id, week=week)
    problems = list_problems(db, factory_id, week=week)
    imports = list_schedule_batches(db, factory_id, week=week, limit=10)
    reports = list_reports(db, factory_id, week=week)
    problem_order_ids = {
        str(problem["inspection_order_id"])
        for problem in problems
        if problem["status"] != "CANCELLED"
    }
    return {
        "factory_id": factory_id,
        "week": week,
        "summary": {
            "total_orders": len(orders),
            "pending_orders": sum(
                1 for item in orders if item["inspection_result"] == "PENDING"
            ),
            "completed_orders": sum(
                1 for item in orders if item["status"] == "COMPLETED"
            ),
            "problem_orders": len(problem_order_ids),
        },
        "orders": orders,
        "problems": problems,
        "imports": imports,
        "reports": reports,
    }


def _rename_group_result_out(result: object) -> dict[str, object]:
    return {
        "group_id": result.group_id,
        "success": bool(result.success),
        "base_name": result.base_name,
        "files": [
            {
                "source_file_name": item.source_file_name,
                "target_file_name": item.target_file_name,
                "source_sha256": item.sha256,
                "size_bytes": item.size_bytes,
            }
            for item in result.files
        ],
        "issues": [
            {
                "code": issue.code,
                "message": issue.message,
                "source_file_name": issue.source_file_name,
            }
            for issue in result.issues
        ],
    }


def rename_batch_to_out(db: Session, batch: QcReportRenameBatch) -> dict[str, object]:
    groups = list(
        db.scalars(
            select(QcReportRenameGroup)
            .where(
                QcReportRenameGroup.factory_id == batch.factory_id,
                QcReportRenameGroup.batch_id == batch.id,
            )
            .order_by(QcReportRenameGroup.client_group_id)
        ).all()
    )
    files_by_group: dict[str, list[QcReportRenameSourceFile]] = defaultdict(list)
    if groups:
        for source in db.scalars(
            select(QcReportRenameSourceFile)
            .where(
                QcReportRenameSourceFile.factory_id == batch.factory_id,
                QcReportRenameSourceFile.group_id.in_([item.id for item in groups]),
            )
            .order_by(
                QcReportRenameSourceFile.group_id,
                QcReportRenameSourceFile.sequence,
                QcReportRenameSourceFile.source_file_name,
            )
        ).all():
            files_by_group[source.group_id].append(source)
    return {
        "id": batch.id,
        "factory_id": batch.factory_id,
        "status": batch.status,
        "rule_version": batch.rule_version,
        "fingerprint": batch.fingerprint,
        "group_count": batch.group_count,
        "source_file_count": batch.source_file_count,
        "source_size_bytes": batch.source_size_bytes,
        "successful_group_count": batch.successful_group_count,
        "failed_group_count": batch.failed_group_count,
        "revision": batch.revision,
        "archive_file_name": batch.archive_file_name,
        "archive_sha256": batch.archive_sha256,
        "archive_size_bytes": batch.archive_size_bytes,
        "created_at": batch.created_at,
        "executed_at": batch.executed_at,
        "groups": [
            {
                "group_id": group.client_group_id,
                "success": group.status == "READY",
                "base_name": group.base_name or None,
                "files": [
                    {
                        "source_file_name": source.source_file_name,
                        "target_file_name": source.target_file_name,
                        "source_sha256": source.source_sha256,
                        "size_bytes": source.size_bytes,
                    }
                    for source in files_by_group[group.id]
                    if group.status == "READY"
                ],
                "issues": [
                    item
                    for item in _json_list(group.issues_json)
                    if isinstance(item, dict)
                ],
            }
            for group in groups
        ],
    }


def preview_rename_batch(
    db: Session,
    *,
    metadata: QcRenamePreviewMeta,
    upload_files: dict[str, tuple[str, bytes]],
    user: AuthContext,
) -> dict[str, object]:
    factory_id = require_qc_factory(metadata.factory_id)
    referenced_ids = {
        source.file_id for group in metadata.groups for source in group.files
    }
    if set(upload_files) != referenced_ids:
        missing = sorted(referenced_ids - set(upload_files))
        extra = sorted(set(upload_files) - referenced_ids)
        detail = []
        if missing:
            detail.append(f"缺少上传文件：{','.join(missing)}")
        if extra:
            detail.append(f"存在未分组上传文件：{','.join(extra)}")
        raise HTTPException(status_code=422, detail="；".join(detail))
    groups: list[ReportRenameGroup] = []
    payload_groups: list[dict[str, object]] = []
    for group in metadata.groups:
        sources: list[ReportSourceFile] = []
        payload_files: list[dict[str, object]] = []
        for source_meta in group.files:
            uploaded_name, content = upload_files[source_meta.file_id]
            if Path(uploaded_name).name != Path(source_meta.source_file_name).name:
                raise HTTPException(
                    status_code=422,
                    detail=f"上传文件 {source_meta.file_id} 的名称与元数据不一致",
                )
            sources.append(
                ReportSourceFile(
                    source_file_name=source_meta.source_file_name,
                    content=content,
                    sequence=source_meta.sequence,
                )
            )
            payload_files.append(
                {
                    "file_id": source_meta.file_id,
                    "source_file_name": source_meta.source_file_name,
                    "sequence": source_meta.sequence,
                    "sha256": hashlib.sha256(content).hexdigest(),
                    "size_bytes": len(content),
                }
            )
        groups.append(
            ReportRenameGroup(
                group_id=group.group_id,
                files=tuple(sources),
                is_caixing=group.is_caixing,
                item_number=group.item_number,
                po=group.customer_po_no,
                actual_inspection_date=group.actual_inspection_date,
                export_country=group.export_country or None,
                report_number=group.report_number or None,
                quantity=group.quantity or None,
            )
        )
        payload_groups.append(
            {
                **group.model_dump(exclude={"files"}),
                "files": payload_files,
            }
        )
    payload_hash = _payload_sha256(
        {
            "factory_id": factory_id,
            "groups": payload_groups,
        }
    )
    replay = _idempotent_response(
        db, user, "RENAME_PREVIEW", metadata.request_id, payload_hash
    )
    if replay is not None:
        return replay
    try:
        result = rename_report_batch(groups)
    except ReportRenameBatchError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    timestamp = now_text()
    batch = QcReportRenameBatch(
        id=f"QRB-{uuid4().hex}",
        factory_id=factory_id,
        status="PREVIEWED",
        rule_version=RENAME_RULE_VERSION,
        fingerprint=result.fingerprint,
        group_count=len(groups),
        source_file_count=len(referenced_ids),
        source_size_bytes=sum(len(item[1]) for item in upload_files.values()),
        successful_group_count=result.successful_group_count,
        failed_group_count=result.failed_group_count,
        preview_json=_canonical_json(
            {
                "groups": [_rename_group_result_out(item) for item in result.group_results]
            }
        ),
        revision=1,
        preview_user_id=user.id,
        preview_request_id=metadata.request_id,
        preview_payload_sha256=payload_hash,
        execute_user_id="",
        execute_request_id="",
        archive_file_name="",
        archive_sha256="",
        archive_size_bytes=0,
        archive_bytes=None,
        created_at=timestamp,
        executed_at="",
    )
    db.add(batch)
    result_by_id = {item.group_id: item for item in result.group_results}
    for group_meta in metadata.groups:
        group_result = result_by_id[group_meta.group_id]
        group_record = QcReportRenameGroup(
            id=f"QRG-{uuid4().hex}",
            factory_id=factory_id,
            batch_id=batch.id,
            client_group_id=group_meta.group_id,
            is_caixing=group_meta.is_caixing,
            export_country=group_meta.export_country,
            report_number=group_meta.report_number,
            item_number=group_meta.item_number,
            customer_po_no=group_meta.customer_po_no,
            quantity=group_meta.quantity,
            actual_inspection_date=group_meta.actual_inspection_date,
            status="READY" if group_result.success else "BLOCKED",
            base_name=group_result.base_name or "",
            issues_json=_canonical_json(
                [
                    {
                        "code": issue.code,
                        "message": issue.message,
                        "source_file_name": issue.source_file_name,
                    }
                    for issue in group_result.issues
                ]
            ),
        )
        db.add(group_record)
        target_by_source = {
            item.source_file_name: item.target_file_name for item in group_result.files
        }
        for source_meta in group_meta.files:
            _, content = upload_files[source_meta.file_id]
            db.add(
                QcReportRenameSourceFile(
                    id=f"QRF-{uuid4().hex}",
                    factory_id=factory_id,
                    group_id=group_record.id,
                    source_file_name=source_meta.source_file_name,
                    source_sha256=hashlib.sha256(content).hexdigest(),
                    size_bytes=len(content),
                    sequence=source_meta.sequence,
                    source_bytes=content,
                    target_file_name=target_by_source.get(source_meta.source_file_name, ""),
                )
            )
    _flush(db, conflict_message="改名预览创建冲突，请刷新后重试")
    response = rename_batch_to_out(db, batch)
    _audit(
        db,
        user,
        factory_id,
        "REPORT_RENAME_PREVIEW_CREATED",
        "qc_report_rename_batch",
        batch.id,
        metadata.request_id,
        after={
            "fingerprint": result.fingerprint,
            "successful_group_count": result.successful_group_count,
            "failed_group_count": result.failed_group_count,
        },
    )
    _record_idempotency(
        db,
        user,
        factory_id,
        "RENAME_PREVIEW",
        metadata.request_id,
        payload_hash,
        response,
    )
    _commit(db, conflict_message="相同 request_id 的改名预览已存在")
    return response


def _validated_rename_archive(batch: QcReportRenameBatch) -> bytes:
    archive = bytes(batch.archive_bytes or b"")
    if batch.status != "EXECUTED" or not archive:
        raise HTTPException(status_code=409, detail="改名批次尚未执行")
    if (
        len(archive) != batch.archive_size_bytes
        or hashlib.sha256(archive).hexdigest() != batch.archive_sha256
    ):
        raise HTTPException(status_code=409, detail="已执行批次的 ZIP 完整性校验失败")
    return archive


def execute_rename_batch(
    db: Session,
    batch_id: str,
    payload: QcRenameExecuteRequest,
    user: AuthContext,
) -> tuple[dict[str, object], bytes]:
    factory_id = require_qc_factory(payload.factory_id)
    payload_data = payload.model_dump(mode="json")
    payload_hash = _payload_sha256(payload_data)
    operation = f"RENAME_EXECUTE:{batch_id}"
    replay = _idempotent_response(db, user, operation, payload.request_id, payload_hash)
    if replay is not None:
        batch = db.get(QcReportRenameBatch, batch_id)
        if batch is None or batch.factory_id != factory_id:
            raise HTTPException(status_code=409, detail="改名执行记录不完整")
        return replay, _validated_rename_archive(batch)
    batch = db.get(QcReportRenameBatch, batch_id)
    if batch is None or batch.factory_id != factory_id:
        raise HTTPException(status_code=404, detail="改名预览批次不存在")
    if batch.status == "EXECUTED":
        archive = _validated_rename_archive(batch)
        # POST execute 同时是可恢复下载入口：即使首次成功响应在网络中丢失，
        # 客户端以新的 request_id 重试也能取得同一份不可变工件。
        return rename_batch_to_out(db, batch), archive
    if batch.revision != payload.expected_revision:
        raise HTTPException(status_code=409, detail="改名预览已变化，请刷新后重试")
    if batch.status != "PREVIEWED":
        raise HTTPException(status_code=409, detail="改名批次已经执行")
    group_records = list(
        db.scalars(
            select(QcReportRenameGroup)
            .where(
                QcReportRenameGroup.factory_id == factory_id,
                QcReportRenameGroup.batch_id == batch.id,
            )
            .order_by(QcReportRenameGroup.client_group_id)
        ).all()
    )
    source_records = list(
        db.scalars(
            select(QcReportRenameSourceFile)
            .where(
                QcReportRenameSourceFile.factory_id == factory_id,
                QcReportRenameSourceFile.group_id.in_([item.id for item in group_records]),
            )
            .order_by(QcReportRenameSourceFile.group_id, QcReportRenameSourceFile.sequence)
        ).all()
    )
    by_group: dict[str, list[QcReportRenameSourceFile]] = defaultdict(list)
    for source in source_records:
        if hashlib.sha256(bytes(source.source_bytes)).hexdigest() != source.source_sha256:
            raise HTTPException(status_code=409, detail="源文件完整性校验失败，禁止执行")
        by_group[source.group_id].append(source)
    groups = [
        ReportRenameGroup(
            group_id=group.client_group_id,
            files=tuple(
                ReportSourceFile(
                    source_file_name=source.source_file_name,
                    content=bytes(source.source_bytes),
                    sequence=source.sequence,
                )
                for source in by_group[group.id]
            ),
            is_caixing=bool(group.is_caixing),
            item_number=group.item_number,
            po=group.customer_po_no,
            actual_inspection_date=group.actual_inspection_date,
            export_country=group.export_country or None,
            report_number=group.report_number or None,
            quantity=group.quantity or None,
        )
        for group in group_records
    ]
    try:
        result = rename_report_batch(groups)
    except ReportRenameBatchError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    if result.fingerprint != batch.fingerprint:
        raise HTTPException(status_code=409, detail="改名预览已失效，源文件或字段发生变化")
    if result.successful_group_count == 0:
        raise HTTPException(status_code=422, detail="没有可执行的报告组，禁止生成空 ZIP")
    archive_hash = hashlib.sha256(result.archive_bytes).hexdigest()
    batch.status = "EXECUTED"
    batch.revision += 1
    batch.execute_user_id = user.id
    batch.execute_request_id = payload.request_id
    batch.archive_file_name = result.archive_file_name
    batch.archive_sha256 = archive_hash
    batch.archive_size_bytes = len(result.archive_bytes)
    batch.archive_bytes = result.archive_bytes
    batch.executed_at = now_text()
    _flush(db, conflict_message="改名执行冲突，请刷新后重试")
    response = rename_batch_to_out(db, batch)
    _audit(
        db,
        user,
        factory_id,
        "REPORT_RENAME_EXECUTED",
        "qc_report_rename_batch",
        batch.id,
        payload.request_id,
        after={
            "archive_file_name": result.archive_file_name,
            "archive_sha256": archive_hash,
            "archive_size_bytes": len(result.archive_bytes),
        },
    )
    _record_idempotency(
        db,
        user,
        factory_id,
        operation,
        payload.request_id,
        payload_hash,
        response,
    )
    _commit(db, conflict_message="改名执行请求冲突，请刷新后重试")
    return response, result.archive_bytes


def get_rename_batch(
    db: Session, factory_id: str, batch_id: str
) -> dict[str, object]:
    factory_id = require_qc_factory(factory_id)
    batch = db.get(QcReportRenameBatch, batch_id)
    if batch is None or batch.factory_id != factory_id:
        raise HTTPException(status_code=404, detail="改名批次不存在")
    return rename_batch_to_out(db, batch)


def get_rename_archive(
    db: Session, factory_id: str, batch_id: str
) -> tuple[str, bytes]:
    factory_id = require_qc_factory(factory_id)
    batch = db.get(QcReportRenameBatch, batch_id)
    if batch is None or batch.factory_id != factory_id:
        raise HTTPException(status_code=404, detail="改名批次不存在")
    return batch.archive_file_name, _validated_rename_archive(batch)
