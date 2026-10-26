"""Conservative, evidence-preserving staging of heterogeneous carton size tables.

This adapter prepares a review workbook. It never creates master records and never
invents product names, inherits an unlabelled SKU, or substitutes shipping sizes.
"""
from io import BytesIO
import math
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.services.carton_master_import import HEADERS

EXTRA_HEADERS = ["待核问题", "原表位置"]
_NUMBER = r"(?:\d+(?:\.\d+)?)"


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value if math.isfinite(value) else None
    match = re.fullmatch(rf"\s*({_NUMBER})\s*(?:cm|mm|inch|in|kg|英寸|厘米|毫米|[\"″])?\s*", text(value), re.I)
    return float(match[1]) if match else None


def packing(value):
    raw = text(value)
    if number(value) is not None and re.fullmatch(_NUMBER, raw):
        return number(value)
    match = re.fullmatch(rf"0\s*/\s*({_NUMBER})", raw)
    if not match:
        match = re.fullmatch(rf"({_NUMBER})\s*(?:PCS\s*/\s*(?:CTN|CNT|箱)|个|只)", raw, re.I)
    return float(match[1]) if match else None


def unit(value):
    raw = text(value).lower()
    if "mm" in raw or "毫米" in raw:
        return "mm"
    if "cm" in raw or "厘米" in raw:
        return "cm"
    if re.search(r"inch|英寸|\bin\b|[\"″]", raw):
        return "inch"
    return ""


def raw_row(row):
    return "；".join(f"{get_column_letter(i + 1)}={text(v)}" for i, v in enumerate(row) if text(v))


def _new(sheet, row_no, row, code, quality, dimensions, size_unit, count, net, gross, note, problems=None):
    problems = list(problems or [])
    code = text(code)
    if not code:
        problems.append("缺货号；未从上一行继承")
    elif not re.fullmatch(r"[#A-Za-z0-9][A-Za-z0-9#_. -]*", code):
        problems.append("货号含组合编号或描述，需确认完整货号及是否拆分")
    if code.startswith("=") or "DISPIMG(" in code.upper():
        problems.append("图片/公式不是货号")
    problems.append("缺独立产品名称，按要求留空")
    if not text(quality):
        problems.append("缺纸质")
    if any(v is None or v <= 0 for v in dimensions):
        problems.append("做箱长宽高缺失或非正数")
    if not size_unit:
        problems.append("做箱尺寸单位待确认")
    if count is None or count <= 0:
        problems.append("每箱装产品数缺失或包装层级待确认")
    parsed_weights = [number(net), number(gross)]
    for i, label in enumerate(("净重", "毛重")):
        value = parsed_weights[i]
        if value is None:
            problems.append(f"缺{label}参考值（选填）" if not text((net, gross)[i]) else f"{label}格式待核")
        elif value < 0 or abs(value - round(value, 4)) > 1e-9:
            problems.append(f"{label}精度或数值不符合系统要求，原值留备注")
            parsed_weights[i] = None
        else:
            parsed_weights[i] = round(value, 4)
    if all(v is not None for v in parsed_weights) and parsed_weights[1] < parsed_weights[0]:
        problems.append("毛重小于净重；保留原值，未自动对调")
    location = f"{sheet}!{row_no}"
    return {"values": [code or None, None, f"{sheet}-R{row_no}", "普通箱", text(quality) or None,
                        *dimensions, size_unit or None, "个", count, "否", "停用", "", *parsed_weights],
            "problems": problems, "notes": [f"来源：{location}", note, f"原始行：{raw_row(row)}"],
            "location": location}


def _dimension_values(values, header_unit="", outer=None):
    problems = []
    dims = [number(v) for v in values]
    explicit = {unit(v) for v in values if unit(v)}
    chosen = next(iter(explicit)) if len(explicit) == 1 else header_unit
    if len(explicit) > 1:
        chosen = ""
        problems.append("同一组三维的单位不一致")
    if outer and all(v and v > 0 for v in dims) and all(number(v) for v in outer):
        ratios = [number(o) / d for o, d in zip(outer, dims)]
        if chosen == "cm" and all(2.3 < ratio < 3.1 for ratio in ratios):
            chosen = ""
            problems.append("表头标CM但箱唛尺寸约为2.54倍，疑似英寸，未擅自纠正")
        elif not chosen and all(2.3 < ratio < 3.1 for ratio in ratios):
            problems.append("数值疑似英寸，仅供核对，单位未自动填入")
    if chosen == "mm":
        dims = [round(v / 10, 8) if v is not None else None for v in dims]
        chosen = "cm"
        problems.append("原尺寸明确为mm，已按10mm=1cm换算")
    return dims, chosen, problems


def _finish(records, source_rows, warnings):
    rows = [HEADERS["configurations"] + EXTRA_HEADERS]
    notes = [["原表位置", "原始文字/完整备注", "整理结果"]]
    identities = {}
    for record in records:
        values = record["values"]
        # Block-row weights may come from a later shipping-size row; validate those too.
        for col, label in ((14, "净重"), (15, "毛重")):
            value = values[col]
            if value is not None and (value < 0 or abs(value - round(value, 4)) > 1e-9):
                record["problems"].append(f"{label}精度或数值不符合系统要求，原值留备注")
                values[col] = None
        if values[14] is not None and values[15] is not None and values[15] < values[14]:
            record["problems"].append("毛重小于净重；保留原值，未自动对调")
        identity = tuple(values[i] for i in (0, 4, 5, 6, 7, 8, 10, 14, 15))
        if identity in identities:
            record["problems"].append(f"与{identities[identity]}尺寸/装箱/重量相同；配卡及版本可能不同，需确认后合并")
        else:
            identities[identity] = record["location"]
        note = "；".join(dict.fromkeys(n for n in record["notes"] if n))
        if len(note) > 1000:
            record["problems"].append("完整备注超过系统1000字上限，需人工精简；完整内容见原表文字索引")
            values[13] = f"来源：{record['location']}；完整配卡/包装备注超过1000字，见原表文字索引同位置；确认精简后才能导入。"
        else:
            values[13] = note
        problems = "；".join(dict.fromkeys(record["problems"]))
        # Staging never infers calculation links; align its review fields after
        # every current import column, including newly added optional inputs.
        rows.append(values + [None] * (len(HEADERS["configurations"]) - len(values)) + [problems, record["location"]])
        notes.append([record["location"], note, problems])
    notes.extend([[location, raw, "原始文字保留（包含表头、配件、未识别或排除行）"] for location, raw in source_rows])
    return {"sheets": [{"name": "货号包装整理", "rows": rows}, {"name": "原表文字索引", "rows": notes}],
            "warnings": warnings + [f"共整理{len(records)}条包装候选；产品名称留空，状态为停用，需补齐及核对后再业务导入。",
            "原表图片、图纸和箱唛版式未转成文字或包装配置，请保留原文件核对。配卡/平卡/内箱说明写入备注，不生成采购数量。",
            "做箱尺寸和箱唛外尺寸分别保留；未标单位、组合货号及多层包装不自动猜测。不同来源、旧版/新版不自动覆盖。"]}


def adapt_carton_sheets(sheets, red_rows=None):
    """Pure staging function; sheets are cached scalar matrices, never instructions."""
    records, source_rows, warnings = [], [], []
    red_rows = red_rows or set()
    for sheet, matrix in sheets.items():
        # Generic headers may sit beyond the fixed customer layouts' 28 columns.
        # Reserve the three derived L/W/H positions after any rightmost header;
        # absent values stay unknown instead of crashing or borrowing another row.
        width = max(28, max((len(row) for row in matrix), default=0) + 3)
        data = [list(row) + [None] * (width - len(row)) for row in matrix]
        source_rows.extend((f"{sheet}!{n}", raw_row(row)) for n, row in enumerate(data, 1) if any(text(v) for v in row))
        if sheet in {"柜型", "Sheet1"} or not any(any(text(v) for v in row) for row in data):
            warnings.append(f"{sheet}：空表或非货号包装资料，未生成包装行；非空文字保留在索引。")
            continue
        if sheet == "9549魔法兔东莞":
            active = None
            for n, row in enumerate(data[3:], 4):
                if text(row[1]) and re.search(r"\d", text(row[1])):
                    match = re.search(r"^(\d+)个.*装1箱", text(row[2]))
                    dims, size_unit, problems = _dimension_values(row[5:8], unit(row[4]))
                    active = _new(sheet, n, row, row[1], row[3], dims, size_unit,
                                  int(match[1]) if match else None, row[9], row[8], text(row[2]), problems)
                    if text(row[1]).startswith("9459"):
                        active["problems"].append("原货号9459与表名9549不同，未擅自改号")
                    records.append(active)
                elif active and any(text(v) for v in row):
                    active["notes"].append(f"配套/续行{n}：{raw_row(row)}")
                    if "装箱尺寸" in text(row[2]) and number(row[8]) is not None and number(row[9]) is not None:
                        active["values"][14:16] = [number(row[9]), number(row[8])]
                        active["problems"] = [p for p in active["problems"] if p not in {"缺净重参考值（选填）", "缺毛重参考值（选填）"}]
            continue
        if sheet == "湖南转-77351":
            active = None
            for n, row in enumerate(data[1:], 2):
                if not any(text(v) for v in row):
                    continue
                code, mixed = text(row[1]), text(row[2])
                accessory = bool(re.search(r"保护卡|护角|平卡|滑板纸|滑片纸|U行卡", code + mixed))
                if not code or accessory:
                    if active:
                        active["notes"].append(f"配套/续行{n}：{raw_row(row)}")
                        if any(word in mixed for word in ("围板", "围版", "天盖", "地盖")):
                            active["problems"].append("组合箱含围板/天地盖，纸品层级与用量需人工整理")
                    continue
                dims, size_unit, problems = _dimension_values(row[3:6], outer=row[6:9])
                quality = None
                count = packing(mixed)
                match = re.fullmatch(r"([A-Za-z][A-Za-z0-9=+ ]*)\s*([（(].*?[）)])\s*(\d+)\s*(?:个|只)?", mixed)
                if match:
                    quality, count = match[1].strip() + match[2], int(match[3])
                if not quality:
                    problems.append("装箱/纸质混写，纸质未自动猜测")
                if re.search(r"装箱数|天地盖|围板|围版|用最新", mixed + code):
                    problems.append("包装组合/版本或装箱文字需确认")
                active = _new(sheet, n, row, code, quality, dims, size_unit, count, row[14], row[15],
                              f"原装箱/纸质：{mixed}", problems)
                records.append(active)
            continue
        if sheet in {"Zuru清溪", "Zuru河源2"}:
            active = None
            for n, row in enumerate(data[1:], 2):
                if not text(row[0]):
                    if active and any(text(v) for v in row):
                        active["notes"].append(f"配套/续行{n}：{raw_row(row)}")
                    continue
                spec = text(row[3])
                part = rf"({_NUMBER})\s*(mm|cm|inch|in|[\"″])?"
                match = re.search(rf"{part}\s*[*x×]\s*{part}\s*[*x×]\s*{part}", spec, re.I)
                dims = [match[i] + (match[i + 1] or "") for i in (1, 3, 5)] if match else [None] * 3
                dims, size_unit, problems = _dimension_values(dims, unit(spec))
                if any(word in spec for word in ("天盖", "围版", "护角", "平卡")):
                    problems.append("组合箱/配件不能按普通整箱直接使用，需核对层级")
                quality_match = re.search(r"\b[A-Z]=[A-Z]\b", spec)
                active = _new(sheet, n, row, row[0], quality_match[0] if quality_match else None,
                              dims, size_unit, packing(row[1]), row[5], row[4], f"原订箱说明：{spec}", problems)
                records.append(active)
            continue
        if sheet == "9251更新-0309":
            for n, row in enumerate(data[3:], 4):
                if not text(row[0]):
                    continue
                records.append(_new(sheet, n, row, row[0], row[1], [None] * 3, "", packing(row[2]), row[10], row[8],
                    "原表只有封箱后印刷装柜尺寸CM，未替代做箱尺寸", ["缺做箱尺寸，箱唛外尺寸已保留原始行"]))
            continue
        if sheet == "迪奇3":
            for n, row in enumerate(data[13:], 14):
                if not any(text(v) for v in row):
                    continue
                for label, positions, count_col, net_col, gross_col in (("单坑", (9, 10, 11), 5, 8, 7), ("双坑", (21, 22, 23), 16, 20, 19)):
                    if not any(text(row[c]) for c in positions) and not text(row[count_col]):
                        continue
                    dims, size_unit, problems = _dimension_values([row[c] for c in positions], "inch")
                    item = _new(sheet, n, row, row[0], None, dims, size_unit, packing(row[count_col]), row[net_col], row[gross_col],
                                f"{label}包装；纸质需核对；C:E为产品/彩盒尺寸，未用作做箱尺寸", problems)
                    item["values"][2] += f"-{label}"
                    records.append(item)
            continue
        header_index = next((i for i, row in enumerate(data[:10]) if "货号" in [text(v) for v in row]), None)
        if header_index is None:
            warnings.append(f"{sheet}：未识别表头，全部非空文字保留在原表文字索引，未猜测列对应。")
            continue
        header = [text(v) for v in data[header_index]]
        code_col = header.index("货号")
        if sheet == "77152":
            pack_col, quality_col, dims_cols, weight_cols, outer_cols, header_unit = 1, 4, [5, 6, 7], [9, 10], [11, 12, 13], "inch"
            following = data[header_index + 1] if header_index + 1 < len(data) else []
            is_subheader = not any(text(v) for v in following) or any(text(v).upper() in {"纸质", "L", "W", "H", "长", "宽", "高"} for v in following)
            start = header_index + (2 if is_subheader else 1)
        elif sheet == "77625":
            pack_col, quality_col, dims_cols, weight_cols, outer_cols, header_unit = 1, 6, [2, 3, 4], [8, 7], [11, 12, 13], ""
            following = data[header_index + 1] if header_index + 1 < len(data) else []
            is_subheader = not any(text(v) for v in following) or any(text(v).upper() in {"纸质", "L", "W", "H", "长", "宽", "高"} for v in following)
            start = header_index + (2 if is_subheader else 1)
        else:
            pack_col = next((i for i, h in enumerate(header) if h in {"装箱数", "装箱量"}), None)
            quality_col = next((i for i, h in enumerate(header) if "做箱SIZE" in h.upper().replace(" ", "") or "做箱SIZE" in h), None)
            if pack_col is None or quality_col is None:
                warnings.append(f"{sheet}：未确认装箱数或做箱列，非空文字保留供人工映射。")
                continue
            dims_cols = [quality_col + 1, quality_col + 2, quality_col + 3]
            weight_cols = [next((i for i, h in enumerate(header) if "净重" in h), -1), next((i for i, h in enumerate(header) if "毛重" in h), -1)]
            outer_start = next((i for i, h in enumerate(header) if "外箱" in h and "CM" in h.upper()), None)
            outer_cols = [outer_start + i for i in range(3)] if outer_start is not None else []
            header_unit = unit(header[quality_col])
            following = data[header_index + 1] if header_index + 1 < len(data) else []
            is_subheader = not any(text(v) for v in following) or any(text(v).upper() in {"纸质", "L", "W", "H", "长", "宽", "高"} for v in following)
            start = header_index + (2 if is_subheader else 1)
        active = None
        for offset, row in enumerate(data[start:], start + 1):
            if not any(text(v) for v in row):
                continue
            code = text(row[code_col])
            if code.startswith("=") or "DISPIMG(" in code.upper():
                warnings.append(f"{sheet}!{offset}：图片公式不能作为货号，已保留索引。")
                continue
            has_dimensions = sum(number(row[c]) is not None for c in dims_cols) >= 2
            if not code and not has_dimensions:
                if active:
                    active["notes"].append(f"配套/续行{offset}：{raw_row(row)}")
                continue
            # Section captions and repeated headers are retained in the source index.
            if code and not re.search(r"\d", code) and not has_dimensions:
                continue
            dims, size_unit, problems = _dimension_values([row[c] for c in dims_cols],
                unit(row[5]) if sheet == "77625" else header_unit, [row[c] for c in outer_cols])
            if (sheet, offset) in red_rows:
                problems.append("原表红字标待确认，未视为已确认资料")
            if sheet == "360客":
                problems.append("表头注明毛重待更改，请核对重量")
            active = _new(sheet, offset, row, code, row[quality_col], dims, size_unit,
                          packing(row[pack_col]), row[weight_cols[0]] if weight_cols[0] >= 0 else None,
                          row[weight_cols[1]] if weight_cols[1] >= 0 else None,
                          f"原装箱：{text(row[pack_col])}；配卡、内箱、彩盒与箱唛说明见原始行列值", problems)
            # Retain readable accessory labels rather than requiring users to decode columns.
            for c, h in enumerate(header):
                if any(word in h for word in ("啤卡", "平卡", "备注", "彩盒", "纸质/配法")) and text(row[c]):
                    active["notes"].insert(1, f"{h}：{text(row[c])}")
            records.append(active)
    return _finish(records, source_rows, warnings)


def adapt_carton_workbook(content, bounds=None):
    """Caller must run the shared archive/cell-budget preflight before this loader."""
    from app.services.spreadsheet_merge import checked_worksheet_bounds, preflight_xlsx
    if bounds is None:
        bounds = preflight_xlsx(content)
    book = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    sheets, red_rows = {}, set()
    try:
        for sheet in book:
            rows = []
            identifier_columns = set()
            row_count, col_count = checked_worksheet_bounds(bounds, sheet._worksheet_path)
            sheet.reset_dimensions()
            if not row_count:
                sheets[sheet.title] = []
                continue
            for n, cells in enumerate(sheet.iter_rows(max_row=row_count, max_col=col_count), 1):
                values = [cell.value for cell in cells]
                for column, cell in enumerate(cells):
                    if n <= 15 and text(cell.value).upper() in {"货号", "产品编号", "SKU NO", "ITEM", "型号", "型號"}:
                        identifier_columns.add(column)
                    if column in identifier_columns and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        raw = text(cell.value)
                        if re.fullmatch(r"0{2,}", cell.number_format or "") and raw.isdigit():
                            raw = raw.zfill(len(cell.number_format))
                        values[column] = raw
                while values and values[-1] is None:
                    values.pop()
                rows.append(values)
                if sheet.title in {"Green toy", "360客", "BJ新客", "建文"}:
                    for cell in cells:
                        color = cell.font.color if getattr(cell, 'has_style', False) else None
                        if cell.value is not None and color and color.type == "rgb" and color.rgb[-6:].upper() in {"FF0000", "C00000"}:
                            red_rows.add((sheet.title, n))
            sheets[sheet.title] = rows
    finally:
        book.close()
    return adapt_carton_sheets(sheets, red_rows)
