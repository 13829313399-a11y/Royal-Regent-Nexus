"""Sewing row semantics shared by calculation, import and export."""
import re
import struct
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def is_sewing_labor(item, part=""):
    return "人工" in f"{item}{part}" or any(str(value).strip() in {"裁床", "车缝", "手工"} for value in (item, part))


def uses_sewing_quoted_price(row):
    currency = str(row.get("unit_price_source_currency") or (
        "HKD" if row.get("unit_price_rmb") in (None, "") and row.get("unit_price_hkd") not in (None, "") else "RMB"
    )).strip().upper()
    return row.get("quoted_price_hkd") not in (None, "") and (
        row.get("usage") in (None, "") or row.get(f"unit_price_{currency.lower()}") in (None, "")
    )


def sewing_formula_multiplier(formula, source_row, cost_column, markup_column):
    """Recognize only same-row cost × markup × numeric literal; never evaluate Excel."""
    if cost_column is None or markup_column is None:
        return None
    cost = f"{get_column_letter(cost_column + 1)}{source_row}"
    markup = f"{get_column_letter(markup_column + 1)}{source_row}"
    value = str(formula or "").upper().replace("$", "").replace(" ", "").lstrip("=")
    match = re.fullmatch(rf"{cost}\*{markup}\*([0-9]+(?:\.[0-9]+)?)", value)
    return Decimal(match[1]) if match else None


def sewing_workbook_formulas(content, sheet_name, max_rows, max_columns):
    """Read bounded formula text, including BIFF8 tokens, without executing it."""
    if not content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        book = load_workbook(BytesIO(content), data_only=False, read_only=True)
        try:
            return {(cell.row, cell.column - 1): cell.value
                    for row in book[sheet_name].iter_rows(max_row=max_rows, max_col=max_columns)
                    for cell in row if cell.data_type == "f"}
        finally:
            book.close()
    import olefile
    import xlrd
    from xlrd.formula import FMLA_TYPE_CELL, decompile_formula

    book = xlrd.open_workbook(file_contents=content, on_demand=True)
    try:
        with olefile.OleFileIO(BytesIO(content)) as ole:
            data = ole.openstream("Workbook" if ole.exists("Workbook") else "Book").read()

        def records(offset):
            while offset + 4 <= len(data):
                code, length = struct.unpack_from("<HH", data, offset)
                payload = data[offset + 4:offset + 4 + length]
                if len(payload) != length:
                    raise ValueError("XLS 公式记录不完整，请重新保存后导入")
                yield code, payload
                offset += length + 4
                if code == 0x000A:
                    break

        offsets = [struct.unpack_from("<I", payload)[0] for code, payload in records(0)
                   if code == 0x0085 and len(payload) >= 8 and payload[5] == 0]
        if book.biff_version != 80 or len(offsets) != book.nsheets:
            raise ValueError("车缝 XLS 需为 Excel 97-2003 格式，请另存为 XLSX 后导入")
        formulas = {}
        for code, payload in records(offsets[book.sheet_names().index(sheet_name)]):
            if code != 0x0006 or len(payload) < 22:
                continue
            row, column = struct.unpack_from("<HH", payload)
            if row >= max_rows or column >= max_columns:
                continue
            size = struct.unpack_from("<H", payload, 20)[0]
            try:
                formulas[(row + 1, column)] = decompile_formula(book, payload[22:22 + size], size, FMLA_TYPE_CELL, row, column)
            except (ValueError, IndexError, AssertionError, NotImplementedError):
                formulas[(row + 1, column)] = "unsupported"
        return formulas
    finally:
        book.release_resources()
