"""Bound Shushupapa imports to the first order section of each sheet."""
from copy import deepcopy
import re

from openpyxl.utils import get_column_letter


_TOTALS = {"合计", "总计", "總計", "小计", "小計"}
_SECTIONS = {"备料", "备料单", "备料订单", "已走货", "已走货订单", "取消单"}
_SUM = re.compile(
    r"^=\s*SUM\s*\(\s*(?P<start>\$?[A-Z]{1,3}\$?)(?P<first>\d+)"
    r"\s*:\s*(?P<end>\$?[A-Z]{1,3}\$?)(?P<last>\d+)\s*\)\s*$",
    re.I,
)


def _label(cell):
    if cell.data_type == "f" or not isinstance(cell.value, str):
        return ""
    return re.sub(r"[\s:：]+", "", cell.value)


def boundary(sheet):
    """Return the first structural footer, excluding order types and remarks."""
    identity_columns = (4, 5, 6, 8) if sheet.title == "ITEM表" else (3, 4, 5, 7)
    candidates = {}
    for (number, column), cell in sheet._cells.items():
        if number < 4:
            continue
        label = _label(cell)
        if column <= 10 and label in _TOTALS:
            candidates[number] = True
        elif column <= 2 and label in _SECTIONS:
            candidates.setdefault(number, False)
    for number, is_total in sorted(candidates.items()):
        # A real order may itself have type 备料单 or a product named 合计.
        if any(
            sheet.cell(number, column).value not in (None, "")
            and sheet.cell(number, column).data_type != "f"
            and _label(sheet.cell(number, column)) not in _TOTALS | _SECTIONS
            for column in identity_columns
        ):
            continue
        return number, is_total
    from app.services.customer_order_unified import CustomerOrderUnifiedError
    raise CustomerOrderUnifiedError(f"{sheet.title} 未找到顶部正单区的合计或分区边界，请核对模板")


def extend_total(sheet, total_row, last_order_row):
    """Include reused/inserted detail rows in the first proven local SUM only."""
    for cell in sheet[total_row]:
        if cell.data_type != "f" or not isinstance(cell.value, str):
            continue
        match = _SUM.fullmatch(cell.value)
        if not match:
            continue
        column = get_column_letter(cell.column)
        if any(match.group(part).replace("$", "").upper() != column for part in ("start", "end")):
            continue
        first, last = int(match.group("first")), int(match.group("last"))
        if first <= 4 <= last < last_order_row < total_row:
            cell.value = (
                cell.value[:match.start("last")] + str(last_order_row)
                + cell.value[match.end("last"):]
            )


def insert_rows(workbook, target, start, amount):
    """Move rule references as well as the shared inserter's cells/ranges."""
    from app.services.huakang_a_order_legacy.unified_schedule import _insert_rows
    from app.services.huaxing_order_legacy.new_order_excel import _rewrite_formula_for_insert

    # An in-memory template may reuse one rule on multiple worksheets. Its
    # unqualified references belong to each owner, so separate those objects.
    owners = {}
    for sheet in workbook:
        collections = [
            *sheet.conditional_formatting._cf_rules.values(),
            sheet.data_validations.dataValidation,
        ]
        for collection in collections:
            for index, rule in enumerate(collection):
                owner = owners.setdefault(id(rule), sheet.title)
                if owner != sheet.title:
                    collection[index] = deepcopy(rule)
    conditional = [
        (sheet, rule, list(rule.formula))
        for sheet in workbook
        for rules in sheet.conditional_formatting._cf_rules.values()
        for rule in rules if rule.formula
    ]
    validations = [
        (sheet, validation, field, getattr(validation, field))
        for sheet in workbook
        for validation in sheet.data_validations.dataValidation
        for field in ("formula1", "formula2")
        if isinstance(getattr(validation, field), str) and getattr(validation, field)
    ]
    _insert_rows(workbook, target, start, amount)

    def rewrite(sheet, formula):
        leading_equals = formula.startswith("=")
        result = _rewrite_formula_for_insert(
            formula if leading_equals else "=" + formula,
            formula_sheet=sheet.title, target_sheet=target.title,
            # Rules are not subtotal cells: an adjacent range must not expand
            # merely because its formatting/validation region moved.
            formula_row=0, insert_row=start, amount=amount,
        )
        return result if leading_equals else result[1:]

    for sheet, rule, formulas in conditional:
        rule.formula = [rewrite(sheet, formula) for formula in formulas]
    for sheet, validation, field, formula in validations:
        setattr(validation, field, rewrite(sheet, formula))
