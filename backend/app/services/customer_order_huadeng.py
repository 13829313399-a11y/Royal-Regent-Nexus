from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from pathlib import Path
import re
from tempfile import TemporaryDirectory
from typing import Any, Callable

import openpyxl
from openpyxl.utils import get_column_letter

from app.services.customer_order_manual import (
    apply_overrides_to_preview,
    apply_overrides_to_records,
    decorate_manual_resolution_policy,
)
from app.services.huadeng_order_legacy import (
    casdon_po_parser,
    casdon_schedule,
    jakks_new_order_writer,
    jakks_po_parser,
    jakks_schedule,
    simba_po_parser,
    simba_schedule,
    spin_master_new_order_writer,
    spin_master_parser,
    spin_po_parser,
    spin_schedule,
)
from app.services.huaxing_order_legacy.new_order_excel import (
    append_records_to_workbook,
    insert_column_records_after_matching_groups,
)


PREVIEW_SCHEMA_VERSION = "customer-order-huadeng-mapped-preview-v1"


class HuadengCustomerOrderError(ValueError):
    """A Huadeng customer workbook could not be mapped safely."""


@dataclass(frozen=True)
class HuadengCustomerMappingSpec:
    code: str
    name: str
    po_extensions: tuple[str, ...]
    schedule_extensions: tuple[str, ...]
    input_template: str
    target_template: str
    rule_summary: str


@dataclass
class PreparedBatch:
    records: list[dict[str, Any]]
    warnings: list[str]
    sheet_name: str
    legacy_rows: list[dict[int, Any]] = field(default_factory=list)
    export_payload: dict[str, Any] = field(default_factory=dict)
    inheritance: dict[str, Any] = field(default_factory=dict)
    header_row: int = 1


HUADENG_CUSTOMER_MAPPINGS: dict[str, HuadengCustomerMappingSpec] = {
    "goliath": HuadengCustomerMappingSpec(
        "goliath", "Goliath", (".pdf",), (".xlsx",),
        "HUADENG_GOLIATH_FORMAL_PO_V1", "HEYUAN_BUSINESS_UNIFIED_REGIONAL_V3",
        "Far East / BV 欧洲版正式 PO 按列提取订单与金额；欧洲版取产品表交期和装箱量，备注日期不同提示复核；缺失外箱和中文品名仅从最新排期唯一历史值继承。",
    ),
    "casdon": HuadengCustomerMappingSpec(
        "casdon", "Casdon", (".pdf", ".xlsx", ".xlsm"), (".xlsx",),
        "HUADENG_CASDON_PO_V1", "HUADENG_CASDON_SCHEDULE_APPEND_V2",
        "按 PO 修订版和字段完整度去重；新单插入同货号最后一行，并按同货号唯一值继承外箱；"
        "外箱缺失或存在多个值时需人工填写/确认，总箱按数量除以外箱；"
        "USD 按 7.75 换算 HKD，验货期为走货期前 7 天并避开周末。",
    ),
    "jakks": HuadengCustomerMappingSpec(
        "jakks", "Jakks", (".pdf", ".xlsx", ".xlsm"), (".xls", ".xlsx"),
        "HUADENG_JAKKS_CONTRACT_V1", "HUADENG_JAKKS_SCHEDULE_APPEND_V2",
        "支持正式 PO 和 SUPPLEMENTARY CONTRACT 补充合同（箱唛资料）；缺失价格留空，取消单拦截，同订单货号查重并继承唯一产品名称。",
    ),
    "simba": HuadengCustomerMappingSpec(
        "simba", "Simba", (".pdf", ".xlsx", ".xlsm"), (".xlsx", ".xlsm"),
        "HUADENG_SIMBA_RELEASE_ORDER_V1", "HUADENG_SIMBA_SCHEDULE_APPEND_V2",
        "同名 PDF 与 WPS Excel 优先采用 Excel；按修订版、文件类型及字段完整度合并，并从排期安全继承客户资料。",
    ),
    "spin": HuadengCustomerMappingSpec(
        "spin", "Spin", (".pdf", ".xlsx", ".xlsm"), (".xlsx",),
        "HUADENG_SPIN_PO_V1", "HUADENG_SPIN_SCHEDULE_APPEND_V2",
        "按 PO 修订版去重，按客户与货号继承排期主数据；USD 按 7.75 换算 HKD，人工排期字段保持空白。",
    ),
    "spin-master": HuadengCustomerMappingSpec(
        "spin-master", "Spin Master", (".pdf", ".xls", ".xlsx", ".xlsm"),
        (".xls", ".xlsx"), "HUADENG_SPIN_MASTER_PO_V1",
        "HUADENG_SPIN_MASTER_SCHEDULE_APPEND_V2",
        "使用独立 SPIN排期 / SPIN总汇模板；按合同、客户 PO、货号、数量、交期及金额组合去重。",
    ),
}


def get_huadeng_customer_mapping(customer_code: str) -> HuadengCustomerMappingSpec:
    try:
        return HUADENG_CUSTOMER_MAPPINGS[customer_code]
    except KeyError as exc:
        raise HuadengCustomerOrderError(f"不支持的华登客户映射：{customer_code}") from exc


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _joined(*values: Any) -> str:
    return "；".join(dict.fromkeys(text for value in values if (text := _text(value))))


def _combined_hash(file_names: list[str], hashes: list[str]) -> str:
    if len(hashes) == 1:
        return hashes[0]
    payload = "\n".join(
        f"{name}:{digest}" for name, digest in zip(file_names, hashes, strict=True)
    )
    return sha256(payload.encode("utf-8")).hexdigest()


def _duplicate_key(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", _text(value).upper())


def _duplicate_quantity(value: Any) -> str:
    try:
        return format(Decimal(str(value)).normalize(), "f")
    except (InvalidOperation, TypeError, ValueError):
        return _text(value)


def _schedule_line_index(
    records: list[dict[str, Any]],
    *,
    order_fields: tuple[str, ...],
    item_fields: tuple[str, ...],
    quantity_field: str = "quantity",
) -> defaultdict[tuple[str, str], set[str]]:
    result: defaultdict[tuple[str, str], set[str]] = defaultdict(set)
    for record in records:
        order_keys = {_duplicate_key(record.get(field)) for field in order_fields}
        item_keys = {_duplicate_key(record.get(field)) for field in item_fields}
        for order_key in order_keys - {""}:
            for item_key in item_keys - {""}:
                result[(order_key, item_key)].add(_duplicate_quantity(record.get(quantity_field)))
    return result


def _mark_record_from_existing_quantities(
    record: dict[str, Any],
    quantities: set[str],
) -> str:
    if not quantities:
        return ""
    quantity = _duplicate_quantity(record.get("quantity"))
    if quantity and quantity in quantities:
        record["_duplicate_existing"] = True
        return "duplicate"
    record.setdefault("flags", []).append({
        "level": "high",
        "code": "existing_quantity_conflict",
        "text": "相同订单及货号已存在排期，但数量不同；按修改/补单阻断",
    })
    return "conflict"


def _mark_records_from_schedule(
    records: list[dict[str, Any]],
    schedule_records: list[dict[str, Any]],
    *,
    record_order_fields: tuple[str, ...],
    record_item_fields: tuple[str, ...],
    schedule_order_fields: tuple[str, ...],
    schedule_item_fields: tuple[str, ...],
    schedule_quantity_field: str = "quantity",
) -> tuple[int, int]:
    index = _schedule_line_index(
        schedule_records,
        order_fields=schedule_order_fields,
        item_fields=schedule_item_fields,
        quantity_field=schedule_quantity_field,
    )
    duplicates = 0
    conflicts = 0
    for record in records:
        quantities: set[str] = set()
        order_keys = {_duplicate_key(record.get(field)) for field in record_order_fields} - {""}
        item_keys = {_duplicate_key(record.get(field)) for field in record_item_fields} - {""}
        for order_key in order_keys:
            for item_key in item_keys:
                quantities.update(index.get((order_key, item_key), set()))
        outcome = _mark_record_from_existing_quantities(record, quantities)
        duplicates += outcome == "duplicate"
        conflicts += outcome == "conflict"
    return duplicates, conflicts


def _safe_output_name(schedule_file_name: str, customer_name: str) -> str:
    stem = re.sub(r"[\\/:*?\"<>|]+", "_", Path(schedule_file_name).stem).strip(" ._")
    suffix = ".xlsm" if Path(schedule_file_name).suffix.lower() == ".xlsm" else ".xlsx"
    return f"{stem or '华登客户排期'}_{customer_name}新单{suffix}"


def _revision(file_name: str) -> int:
    name = re.sub(r"\(\d+\)\s*$", "", Path(file_name).stem).strip()
    best = 0
    for pattern in (r"rev\.?\s*(\d+)", r"\br\.?\s*(\d+)", r"\bv\.?\s*(\d+)"):
        for match in re.finditer(pattern, name, re.I):
            best = max(best, int(match.group(1)))
    return 1 if best == 0 and re.search(r"\brev\b", name, re.I) else best


def _dedupe_revision_orders(
    orders: list[dict[str, Any]], *, quality_fields: tuple[str, ...] = (),
) -> tuple[list[dict[str, Any]], list[str]]:
    grouped: defaultdict[str, list[tuple[int, tuple[int, int], str, dict[str, Any]]]] = defaultdict(list)
    result: list[dict[str, Any]] = []
    report: list[str] = []
    for order in orders:
        po_number = _text(order.get("po_number"))
        if not po_number:
            result.append(order)
            continue
        header_score = sum(bool(order.get(field)) for field in quality_fields)
        line_score = sum(
            sum(bool(line.get(field)) for field in ("item_code", "material_number", "qty", "unit_price", "total_usd"))
            for line in order.get("lines", [])
        )
        grouped[po_number.upper()].append(
            (_revision(_text(order.get("filename"))), (header_score, line_score), _text(order.get("filename")), order)
        )
    for po_number, entries in grouped.items():
        entries.sort(key=lambda item: (item[0], item[1], item[2]), reverse=True)
        result.append(entries[0][3])
        if len(entries) > 1:
            report.append(
                f"PO {po_number}：保留 {entries[0][2]}，去掉 "
                + "、".join(item[2] for item in entries[1:])
            )
    return result, report


def _excel_serial_text(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, (int, float)):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).date().isoformat()
        except (OverflowError, ValueError):
            pass
    return _text(value)


def _schedule_evidence(file_name: str, sheet_name: str, customers: list[Any]) -> str:
    return " ".join([file_name, sheet_name, *[_text(item) for item in customers]]).upper()


def _casdon_schedule_context(schedule_path: Path) -> tuple[Any, str, int]:
    workbook = openpyxl.load_workbook(schedule_path, read_only=True, data_only=True)
    try:
        worksheet, header_row = casdon_schedule._select_master_sheet(workbook)
        customers = [
            worksheet.cell(row_no, casdon_schedule.COL["customer"]).value
            for row_no in range(header_row + 1, min(worksheet.max_row, header_row + 200) + 1)
        ]
        if "CASDON" not in _schedule_evidence(schedule_path.name, worksheet.title, customers):
            raise HuadengCustomerOrderError("当前排期未识别到 Casdon 客户标识，已拦截跨客户模板")
        return casdon_schedule.build_index(str(schedule_path)), worksheet.title, header_row
    finally:
        workbook.close()


def _spin_schedule_context(schedule_path: Path) -> tuple[Any, str, int]:
    workbook = openpyxl.load_workbook(schedule_path, read_only=True, data_only=True)
    try:
        if spin_schedule.MASTER_SHEET not in workbook.sheetnames:
            raise HuadengCustomerOrderError(
                f"Spin 排期缺少工作表：{spin_schedule.MASTER_SHEET}"
            )
        worksheet = workbook[spin_schedule.MASTER_SHEET]
        customers = [
            worksheet.cell(row_no, spin_schedule.COL["customer"]).value
            for row_no in range(2, min(worksheet.max_row, 201) + 1)
        ]
        if "SPIN" not in _schedule_evidence(schedule_path.name, worksheet.title, customers):
            raise HuadengCustomerOrderError("当前排期未识别到 Spin 客户标识，已拦截跨客户模板")
        return spin_schedule.build_index(str(schedule_path)), worksheet.title, 1
    finally:
        workbook.close()


def _inheritance(index: Any) -> dict[str, Any]:
    return {
        "rule": "same_schedule_same_item_unique_name",
        "mapped_items": sum(
            1 for catalog in index.catalog_by_item.values()
            if catalog.get("cn_name") or catalog.get("english_name")
        ),
        "applied_rows": 0,
        "conflicts": index.inheritance_conflicts,
    }


def _legacy_row_record(
    customer_code: str, order: dict[str, Any], line: dict[str, Any], values: dict[int, Any],
) -> dict[str, Any]:
    schedule = casdon_schedule if customer_code == "casdon" else spin_schedule
    col = schedule.COL
    outer = values.get(col["outer"], "")
    quantity = values.get(col["qty"], "")
    cartons: Any = ""
    try:
        if outer not in (None, "", 0) and quantity not in (None, ""):
            calculated = float(quantity) / float(outer)
            cartons = int(calculated) if calculated.is_integer() else round(calculated, 2)
    except (TypeError, ValueError, ZeroDivisionError):
        pass
    inspection_key = "inspection_date"
    record = {
        "po_no": values.get(col["customer_po"], ""),
        "contract_no": values.get(col["contract"], ""),
        "customer_name": values.get(col["customer"], ""),
        "country": values.get(col["ship_country"], ""),
        "product_no": values.get(col["item"], ""),
        "product_name_zh": values.get(col["cn_name"], ""),
        "product_name_en": values.get(col["english_name"], ""),
        "quantity": quantity,
        "units_per_carton": outer,
        "carton_count": cartons,
        "standard": values.get(col["version"], ""),
        "unit_price_hkd": values.get(col["unit_hkd"], ""),
        "amount_hkd": values.get(col["total_hkd"], ""),
        "packaging": _joined(
            values.get(col["special_note"], ""), values.get(col["carton_mark"], ""),
            values.get(col["customer_label"], ""),
        ),
        "line_q": _excel_serial_text(values.get(col[inspection_key], "")),
        "customer_q": "",
        "requested_ship_date": _excel_serial_text(values.get(col["ship_date"], "")),
        "_source_po_file_name": order.get("filename"),
        "_legacy_values": values,
    }
    if line.get("is_charge"):
        record["standard"] = "附加费用"
    return record


def _prepare_casdon_or_spin(
    customer_code: str, po_files: list[tuple[str, bytes]],
    schedule_file_name: str, schedule_content: bytes, received_date: str,
) -> PreparedBatch:
    is_casdon = customer_code == "casdon"
    parser_class = casdon_po_parser.CasdonPOParser if is_casdon else spin_po_parser.SpinPOParser
    validator = casdon_po_parser.validate if is_casdon else spin_po_parser.validate
    schedule_module = casdon_schedule if is_casdon else spin_schedule
    with TemporaryDirectory(prefix=f"huadeng-{customer_code}-") as temp_dir:
        schedule_path = Path(temp_dir) / Path(schedule_file_name).name
        schedule_path.write_bytes(schedule_content)
        try:
            index, sheet_name, header_row = (
                _casdon_schedule_context(schedule_path)
                if is_casdon else _spin_schedule_context(schedule_path)
            )
        except HuadengCustomerOrderError:
            raise
        except Exception as exc:
            raise HuadengCustomerOrderError(
                f"{schedule_file_name}：无法读取 {get_huadeng_customer_mapping(customer_code).name} 排期：{exc}"
            ) from exc

        orders: list[dict[str, Any]] = []
        warnings: list[str] = []
        for file_name, content in po_files:
            safe_name = re.sub(r"[\\/:*?\"<>|]+", "_", Path(file_name).name) or "po"
            po_path = Path(temp_dir) / f"{len(orders)}-{safe_name}"
            po_path.write_bytes(content)
            try:
                parsed = parser_class().parse(str(po_path))
            except Exception as exc:
                warnings.append(f"{file_name}：解析失败（{exc}），本文件未进入新单")
                continue
            parsed["filename"] = file_name
            if is_casdon:
                parsed["email_received_date"] = received_date
            warnings.extend(validator(parsed, file_name))
            orders.append(parsed)
        if not orders:
            raise HuadengCustomerOrderError("本批 PO 均未识别到可生成的新单明细")
        orders, dedupe = _dedupe_revision_orders(
            orders,
            quality_fields=(
                "po_number", "po_date", "ship_date", "your_reference", "customer",
                "customer_po_header",
            ) if is_casdon else (),
        )
        warnings.extend(dedupe)
        legacy_rows: list[dict[int, Any]] = []
        records: list[dict[str, Any]] = []
        inherited_rows = 0
        existing_duplicates = 0
        existing_conflicts = 0
        for order in orders:
            for line in order.get("lines", []):
                values = schedule_module._compose_row(order, line, index)
                legacy_rows.append(values)
                record = _legacy_row_record(customer_code, order, line, values)
                if is_casdon and not line.get("is_charge") and not values.get(
                    schedule_module.COL["outer"]
                ):
                    item = casdon_schedule.item_base(values.get(schedule_module.COL["item"]))
                    outer_values = index.outer_values_by_item.get(item, ())
                    if len(outer_values) > 1:
                        options = " / ".join(_text(value) for value in outer_values)
                        code = "ambiguous_units_per_carton"
                        message = (
                            f"当前排期货号 {item} 存在多个外箱装箱数（{options}），"
                            "请人工填写本单外箱装箱数"
                        )
                    else:
                        code = "missing_units_per_carton"
                        message = (
                            f"当前排期未找到货号 {item} 的外箱装箱数，"
                            "请人工填写本单外箱装箱数"
                        )
                    record.setdefault("flags", []).append({
                        "level": "high",
                        "code": code,
                        "field": "units_per_carton",
                        "text": message,
                    })
                    warnings.append(f"{order.get('filename')} / {item}：{message}")
                if is_casdon:
                    schedule_key = (
                        casdon_schedule.normalize_po(values.get(schedule_module.COL["contract"])),
                        casdon_schedule.item_base(values.get(schedule_module.COL["item"])),
                    )
                    existing_rows = index.index.get(schedule_key, []) if all(schedule_key) else []
                else:
                    schedule_key = spin_schedule.normalize_contract(
                        values.get(schedule_module.COL["contract"])
                    )
                    existing_rows = index.index.get(schedule_key, []) if schedule_key else []
                outcome = _mark_record_from_existing_quantities(
                    record,
                    {_duplicate_quantity(existing[-1]) for existing in existing_rows},
                )
                existing_duplicates += outcome == "duplicate"
                existing_conflicts += outcome == "conflict"
                records.append(record)
                if values.get(schedule_module.COL["cn_name"]) or values.get(schedule_module.COL["english_name"]):
                    inherited_rows += 1
                if not line.get("is_charge") and not values.get(schedule_module.COL["cn_name"]):
                    warnings.append(
                        f"{order.get('filename')} / {line.get('item_code') or line.get('item_key')}："
                        "当前排期未找到唯一中文品名，已留空"
                    )
                if not is_casdon and not values.get(schedule_module.COL["outer"]):
                    warnings.append(
                        f"{order.get('filename')} / {values.get(schedule_module.COL['item'], '')}："
                        "未识别到外箱装箱数，总箱数留空"
                    )
        inheritance = _inheritance(index)
        inheritance["applied_rows"] = inherited_rows
        if inheritance["conflicts"]:
            warnings.append(
                f"排期中有 {len(inheritance['conflicts'])} 组货号/品名冲突，相关品名未自动继承。"
            )
        if existing_duplicates:
            warnings.append(
                f"当前 {get_huadeng_customer_mapping(customer_code).name} 排期已有 "
                f"{existing_duplicates} 行相同订单，测试阶段需逐项确认。"
            )
        if existing_conflicts:
            warnings.append(
                f"当前排期有 {existing_conflicts} 行相同订单但数量不同，已按修改/补单阻断。"
            )
        return PreparedBatch(
            records,
            warnings,
            sheet_name,
            legacy_rows,
            inheritance=inheritance,
            header_row=header_row,
        )


def _jakks_item_key(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", _text(value).upper())


def _jakks_text_key(value: Any) -> str:
    return re.sub(r"[^A-Z0-9\u4e00-\u9fff]", "", _text(value).upper())


def _apply_jakks_schedule_data(order: dict[str, Any], dataset: dict[str, Any]) -> None:
    candidates: dict[str, dict[str, str]] = {}
    date_code_candidates: dict[str, set[str]] = defaultdict(set)
    records = dataset.get("records") or []
    for record in records:
        key = _jakks_item_key(record.get("sku"))
        name = _text(record.get("product"))
        if key and name:
            candidates.setdefault(key, {})[re.sub(r"\s+", "", name).casefold()] = name
        date_code = _text(record.get("date_code"))
        if key and date_code:
            date_code_candidates[key].add(date_code)
    warnings = order.setdefault("warnings", [])
    for line in order.get("lines") or []:
        names = candidates.get(_jakks_item_key(line.get("item_no")), {})
        if len(names) == 1:
            line["product_name"] = next(iter(names.values()))
            line["product_name_source"] = "最新Jakks排期按货号唯一继承"
        elif len(names) > 1:
            line["product_name_source"] = "PO原始品名（排期同货号多品名）"
            warnings.append(f"{line.get('item_no')}：当前 Jakks 排期存在多个产品名称，未自动猜测。")
        else:
            line["product_name_source"] = "PO原始品名（当前排期无此货号）"
            warnings.append(f"{line.get('item_no')}：当前 Jakks 排期没有该货号，保留 PO 品名。")
        date_codes = date_code_candidates.get(_jakks_item_key(line.get("item_no")), set())
        if not line.get("date_code") and len(date_codes) == 1:
            line["date_code"] = next(iter(date_codes))
        flags = line.setdefault("flags", [])
        if not re.search(r"[\u3400-\u9fff]", _text(line.get("product_name"))):
            flags.append({
                "level": "high",
                "code": "missing_product_name_zh",
                "field": "product_name_zh",
                "text": "当前排期未找到完整中文品名，请人工补录后导出",
            })
        if not _text(line.get("date_code")):
            flags.append({
                "level": "high",
                "code": "missing_date_code",
                "field": "date_code",
                "text": "当前排期未找到可唯一继承的日期码，请人工补录后导出",
            })
    if _text(order.get("contact")):
        return
    order_contract = _jakks_text_key(order.get("contract_no"))
    order_po = _jakks_text_key(order.get("customer_po"))
    contacts = {
        _text(record.get("contact")) for record in records
        if _text(record.get("contact")) and (
            (order_contract and _jakks_text_key(record.get("contract_no")) == order_contract)
            or (order_po and _jakks_text_key(record.get("customer_po")) == order_po)
        )
    }
    if len(contacts) == 1:
        order["contact"] = next(iter(contacts))
        warnings.append(f"合同联系人：{order['contact']}（按合同号/客户 PO 唯一继承）。")
    else:
        warnings.append("PO 未提供且排期无法唯一确定合同联系人，已留空。")


def _dedupe_text(value: Any) -> str:
    return re.sub(r"\s+", " ", _text(value)).casefold()


def _jakks_line_identity(line: dict[str, Any], *, business: bool = False) -> tuple[str, ...]:
    fields = (
        "contract_no", "customer_po", "confirmation_no", "item_no", "quantity", "unit", "ship_date",
    ) if business else (
        "contract_no", "customer_po", "confirmation_no", "item_no", "quantity", "unit", "ship_date",
        "unit_price_usd", "total_usd", "outer_pack", "cartons", "special_note", "po_description",
    )
    return tuple(_dedupe_text(line.get(field)) for field in fields)


def _prepare_jakks(
    po_files: list[tuple[str, bytes]], schedule_file_name: str, schedule_content: bytes,
) -> PreparedBatch:
    with TemporaryDirectory(prefix="huadeng-jakks-") as temp_dir:
        schedule_path = Path(temp_dir) / Path(schedule_file_name).name
        schedule_path.write_bytes(schedule_content)
        try:
            dataset = jakks_schedule.load_dataset(schedule_path)
        except Exception as exc:
            raise HuadengCustomerOrderError(f"{schedule_file_name}：无法读取 Jakks 排期：{exc}") from exc
        parsed_orders: list[dict[str, Any]] = []
        warnings: list[str] = []
        for index, (file_name, content) in enumerate(po_files):
            if re.search(r"(?:^|[-_])CXL(?:[-_]|$)", Path(file_name).stem, re.I):
                warnings.append(f"{file_name}：文件名含 CXL，按取消单拦截。")
                continue
            po_path = Path(temp_dir) / f"po-{index}_{Path(file_name).name}"
            po_path.write_bytes(content)
            try:
                order = jakks_po_parser.parse_po(po_path)
            except Exception as exc:
                warnings.append(f"{file_name}：{exc}")
                continue
            order["filename"] = file_name
            _apply_jakks_schedule_data(order, dataset)
            parsed_orders.append(order)
        combined: list[dict[str, Any]] = []
        seen_exact: set[tuple[str, ...]] = set()
        seen_business: dict[tuple[str, ...], dict[str, Any]] = {}
        duplicate_count = 0
        for order in parsed_orders:
            warnings.extend(f"{order.get('filename')}：{item}" for item in order.get("warnings") or [])
            for raw_line in order.get("lines") or []:
                line = dict(raw_line)
                for field_name in (
                    "order_date", "contact", "customer_po", "confirmation_no", "contract_no",
                    "customer", "country", "ship_date",
                ):
                    if not line.get(field_name):
                        line[field_name] = order.get(field_name)
                line["source_file"] = order.get("filename") or ""
                exact = _jakks_line_identity(line)
                if exact in seen_exact:
                    duplicate_count += 1
                    continue
                seen_exact.add(exact)
                business = _jakks_line_identity(line, business=True)
                if business in seen_business:
                    kept = seen_business[business]
                    warnings.append(
                        "同一 Jakks 订单行出现不同版本，已保留先上传版本并拦截重复："
                        f"{kept.get('source_file')} / {line.get('source_file')} / {line.get('item_no')}。"
                    )
                    continue
                seen_business[business] = line
                combined.append(line)
        if duplicate_count:
            warnings.append(f"批内发现 {duplicate_count} 行完全重复 Jakks 明细，已去重。")
        if not combined:
            raise HuadengCustomerOrderError("本批没有可生成的 Jakks 新单明细；修改/取消单或异常文件已按提示拦截")
        existing_count, conflict_count = _mark_records_from_schedule(
            combined,
            dataset.get("records") or [],
            record_order_fields=("contract_no", "customer_po"),
            record_item_fields=("item_no",),
            schedule_order_fields=("contract_no", "customer_po"),
            schedule_item_fields=("sku",),
        )
        if existing_count:
            warnings.append(f"当前 Jakks 排期已有 {existing_count} 行相同订单，测试阶段需逐项确认。")
        if conflict_count:
            warnings.append(f"当前 Jakks 排期有 {conflict_count} 行相同订单但数量不同，已按修改/补单阻断。")
        return PreparedBatch(
            combined, warnings, "26-Jakks排货表总 Ai",
            export_payload={"filename": f"Jakks批量上传（{len(parsed_orders)}份新单）", "lines": combined},
        )


def _simba_normalize(value: Any) -> str:
    text = _text(value).upper()
    if text.endswith(".0"):
        text = text[:-2]
    return re.sub(r"[^A-Z0-9]", "", text)


def _simba_identity(row: dict[str, Any]) -> tuple[str, ...]:
    contract = _simba_normalize(row.get("contract_no"))
    customer_po = _simba_normalize(row.get("customer_po"))
    item = _simba_normalize(row.get("item_no"))
    if contract or customer_po:
        return (
            "order", contract, customer_po, item,
            _simba_normalize(row.get("master_contract_no")),
            _simba_normalize(row.get("po_contract_no")),
        )
    return ("fallback", item, _simba_normalize(row.get("product_name")), _simba_normalize(row.get("quantity")))


def _simba_signature(row: dict[str, Any]) -> tuple[str, ...]:
    return tuple(
        _simba_normalize(row.get(field_name))
        for field_name in (
            "contract_no", "customer_po", "item_no", "product_name", "quantity", "unit",
            "po_ship_date", "master_contract_no", "po_contract_no",
        )
    )


def _simba_priority(file_name: str, row: dict[str, Any]) -> tuple[int, int]:
    revision = 100 if re.search(
        r"(?i)(?:^|[\s_-])(?:rev(?:ise|ision)?\d*|reduce[\s_-]*qty|修改|修订|更改)(?:$|[\s_-])",
        Path(file_name).stem,
    ) else 0
    file_type = 20 if Path(file_name).suffix.lower() in {".xlsx", ".xlsm"} else 10
    completeness = sum(
        row.get(field_name) not in (None, "")
        for field_name in ("contract_no", "customer_po", "item_no", "product_name", "quantity")
    )
    return revision + file_type, completeness


def _simba_first_english_sentence(value: Any) -> str:
    text = re.sub(r"\s+", " ", _text(value)).strip()
    if not text:
        return ""
    return re.split(r"(?<=[.!?])\s+", text, maxsplit=1)[0].strip()


def _prepare_simba(
    po_files: list[tuple[str, bytes]], schedule_file_name: str, schedule_content: bytes,
    received_date: str,
) -> PreparedBatch:
    try:
        schedule = simba_schedule.read_schedule(schedule_content, filename=schedule_file_name)
    except Exception as exc:
        raise HuadengCustomerOrderError(f"{schedule_file_name}：无法读取 Simba 排期：{exc}") from exc
    excel_stems = {
        Path(name).stem.casefold() for name, _ in po_files
        if Path(name).suffix.lower() in {".xlsx", ".xlsm"}
    }
    selected: list[tuple[str, bytes]] = []
    warnings: list[str] = []
    seen_hashes: set[str] = set()
    paired_pdf_count = 0
    for file_name, content in po_files:
        digest = sha256(content).hexdigest()
        if digest in seen_hashes:
            warnings.append(f"{file_name}：完全相同的重复文件未再次处理。")
            continue
        seen_hashes.add(digest)
        if Path(file_name).suffix.lower() == ".pdf" and Path(file_name).stem.casefold() in excel_stems:
            paired_pdf_count += 1
            continue
        selected.append((file_name, content))
    if paired_pdf_count:
        warnings.append(f"检测到 {paired_pdf_count} 组同名 PDF + WPS Excel，已优先采用 Excel。")

    parsed_files: list[dict[str, Any]] = []
    for file_name, content in selected:
        try:
            parsed = simba_po_parser.parse_po_file(content, filename=file_name)
        except Exception as exc:
            warnings.append(f"{file_name}：{exc}")
            continue
        parsed_files.append(parsed)
        warnings.extend(f"{file_name}：{item}" for item in parsed.get("warnings") or [])
    merged: dict[tuple[str, ...], dict[str, Any]] = {}
    order: list[tuple[str, ...]] = []
    for parsed in parsed_files:
        file_name = _text(parsed.get("filename")) or "PO"
        for source in parsed.get("rows") or []:
            row = dict(source)
            # User-confirmed Simba output rule: PO carton dimensions must not
            # populate the schedule's AK–AM dimension cells.
            for field_name in ("outer_length_cm", "outer_width_cm", "outer_height_cm"):
                row.pop(field_name, None)
            row["english_name"] = _simba_first_english_sentence(row.get("english_name"))
            row["source_file"] = file_name
            row["_source_po_file_name"] = file_name
            row["order_date"] = received_date
            simba_schedule.add_derived_fields(row)
            key = _simba_identity(row)
            current = merged.get(key)
            if current is None:
                merged[key] = row
                order.append(key)
                continue
            if _simba_signature(current) == _simba_signature(row):
                warnings.append(f"{file_name}：同单重复明细未再次写入（{row.get('item_no') or '-'}）。")
                continue
            current_name = _text(current.get("source_file")) or "PO"
            if _simba_priority(file_name, row) > _simba_priority(current_name, current):
                merged[key] = row
                warnings.append(f"同单存在版本差异，采用 {file_name} 替代 {current_name}。")
            else:
                warnings.append(f"同单存在版本差异，保留 {current_name}，未重复写入 {file_name}。")
    records = [merged[key] for key in order]
    if not records:
        detail = "；".join(dict.fromkeys(warnings))
        message = "本批文件未识别到可生成的 Simba 新单明细"
        if detail:
            message += f"。文件处理结果：{detail}"
        raise HuadengCustomerOrderError(message)
    inheritance = simba_schedule.enrich_rows_from_schedule(records, schedule.get("records") or [])
    warnings.append(
        f"已从当前 Simba 排期安全继承 {inheritance['product_names_applied']} 个中文品名、"
        f"{inheritance['exact_fields_applied']} 个合同级字段。"
    )
    if inheritance["product_name_conflicts"]:
        warnings.append(
            "当前 Simba 排期同货号存在多个品名，未自动猜测："
            + "、".join(inheritance["product_name_conflicts"])
        )
    existing_count, conflict_count = _mark_records_from_schedule(
        records,
        schedule.get("records") or [],
        record_order_fields=("contract_no", "customer_po"),
        record_item_fields=("item_no",),
        schedule_order_fields=("contract_no", "customer_po"),
        schedule_item_fields=("item_no",),
    )
    if existing_count:
        warnings.append(f"当前 Simba 排期已有 {existing_count} 行相同订单，测试阶段需逐项确认。")
    if conflict_count:
        warnings.append(f"当前 Simba 排期有 {conflict_count} 行相同订单但数量不同，已按修改/补单阻断。")
    return PreparedBatch(records, warnings, _text(schedule.get("sheet")) or "Simba排期")


def _prepare_spin_master(
    po_files: list[tuple[str, bytes]], schedule_file_name: str, schedule_content: bytes,
) -> PreparedBatch:
    with TemporaryDirectory(prefix="huadeng-spin-master-") as temp_dir:
        schedule_path = Path(temp_dir) / Path(schedule_file_name).name
        schedule_path.write_bytes(schedule_content)
        try:
            schedule = spin_master_parser.validate_schedule_file(schedule_path)
        except Exception as exc:
            raise HuadengCustomerOrderError(f"{schedule_file_name}：无法读取 Spin Master 排期：{exc}") from exc
        records: list[dict[str, Any]] = []
        warnings: list[str] = []
        seen: set[tuple[str, ...]] = set()
        duplicate_count = 0
        for index, (file_name, content) in enumerate(po_files):
            po_path = Path(temp_dir) / f"po-{index}{Path(file_name).suffix.lower()}"
            po_path.write_bytes(content)
            try:
                parsed = spin_master_parser.parse_po_file(po_path)
            except Exception as exc:
                warnings.append(f"{file_name}：{exc}")
                continue
            parsed["source"] = file_name
            for record in spin_master_new_order_writer._records(parsed):
                record["_source_po_file_name"] = file_name
                key = tuple(
                    _text(record.get(field_name)).upper()
                    for field_name in (
                        "contract_no", "customer_po", "item_no", "quantity", "ship_date",
                        "unit_price_usd", "total_usd",
                    )
                )
                if key in seen:
                    duplicate_count += 1
                    continue
                seen.add(key)
                records.append(record)
        if duplicate_count:
            warnings.append(f"本批去除 {duplicate_count} 行完全重复的 Spin Master 明细。")
        if not records:
            raise HuadengCustomerOrderError("本批文件未识别到可生成的 Spin Master 新单明细")
        existing_count, conflict_count = _mark_records_from_schedule(
            records,
            schedule.get("rows") or [],
            record_order_fields=("contract_no", "customer_po"),
            record_item_fields=("item_no",),
            schedule_order_fields=("contract_no", "customer_po"),
            schedule_item_fields=("item_si", "item_no"),
        )
        if existing_count:
            warnings.append(f"当前 Spin Master 排期已有 {existing_count} 行相同订单，测试阶段需逐项确认。")
        if conflict_count:
            warnings.append(f"当前 Spin Master 排期有 {conflict_count} 行相同订单但数量不同，已按修改/补单阻断。")
        return PreparedBatch(records, warnings, _text(schedule.get("sheet")) or "SPIN排期")


def _prepare_batch(
    customer_code: str, po_files: list[tuple[str, bytes]], schedule_file_name: str,
    schedule_content: bytes, received_date: str,
) -> PreparedBatch:
    if customer_code in {"casdon", "spin"}:
        return _prepare_casdon_or_spin(
            customer_code, po_files, schedule_file_name, schedule_content, received_date,
        )
    handlers: dict[str, Callable[[], PreparedBatch]] = {
        "jakks": lambda: _prepare_jakks(po_files, schedule_file_name, schedule_content),
        "simba": lambda: _prepare_simba(
            po_files, schedule_file_name, schedule_content, received_date,
        ),
        "spin-master": lambda: _prepare_spin_master(
            po_files, schedule_file_name, schedule_content,
        ),
    }
    try:
        return handlers[customer_code]()
    except HuadengCustomerOrderError:
        raise
    except Exception as exc:
        raise HuadengCustomerOrderError(str(exc)) from exc


def _record_fields(customer_code: str, record: dict[str, Any]) -> dict[str, Any]:
    if customer_code in {"casdon", "spin"}:
        return record
    if customer_code == "jakks":
        unit_usd = float(record.get("unit_price_usd") or 0)
        quantity = float(record.get("quantity") or 0)
        return {
            "po_no": record.get("customer_po"),
            "contract_no": record.get("contract_no"),
            "customer_name": record.get("customer"),
            "country": record.get("country"),
            "product_no": record.get("item_no"),
            "product_name_zh": record.get("product_name"),
            "product_name_en": record.get("po_description"),
            "quantity": record.get("quantity"),
            "units_per_carton": record.get("outer_pack"),
            "carton_count": record.get("cartons"),
            "standard": _joined(record.get("unit"), record.get("confirmation_no")),
            "unit_price_hkd": round(unit_usd * 7.75, 6) if unit_usd else "",
            "amount_hkd": round(quantity * unit_usd * 7.75, 2) if unit_usd else "",
            "packaging": record.get("special_note"),
            "line_q": record.get("inspection_date"),
            "customer_q": "",
            "requested_ship_date": record.get("ship_date"),
        }
    if customer_code == "simba":
        return {
            "po_no": record.get("customer_po"),
            "contract_no": record.get("contract_no"),
            "customer_name": record.get("customer"),
            "country": record.get("ship_country"),
            "product_no": record.get("item_no"),
            "product_name_zh": record.get("product_name"),
            "product_name_en": record.get("english_name"),
            "quantity": record.get("quantity"),
            "units_per_carton": record.get("outer_pack"),
            "carton_count": record.get("cartons"),
            "standard": _joined(record.get("brand"), record.get("certificate")),
            "unit_price_hkd": record.get("unit_price_hkd"),
            "amount_hkd": record.get("total_hkd"),
            "packaging": _joined(record.get("shipping_mark"), record.get("customer_label"), record.get("special_remark")),
            "line_q": record.get("inspection_date"),
            "customer_q": "",
            "requested_ship_date": record.get("po_ship_date"),
        }
    return {
        "po_no": record.get("customer_po"),
        "contract_no": record.get("contract_no"),
        "customer_name": record.get("customer"),
        "country": "",
        "product_no": record.get("item_no"),
        "product_name_zh": record.get("product_name"),
        "product_name_en": record.get("english_name"),
        "quantity": record.get("quantity"),
        "units_per_carton": record.get("outer_pack"),
        "carton_count": record.get("cartons"),
        "standard": record.get("version"),
        "unit_price_hkd": record.get("unit_price_hkd"),
        "amount_hkd": record.get("total_hkd"),
        "packaging": record.get("special_notes"),
        "line_q": record.get("inspection_date"),
        "customer_q": "",
        "requested_ship_date": record.get("ship_date"),
    }


def _issues(record: dict[str, Any], row_id: str) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    for index, flag in enumerate(record.get("flags") or []):
        severity = "blocked" if flag.get("level") == "high" else "warning"
        code = _text(flag.get("code")) or f"risk_{index + 1}"
        issues.append({
            "severity": severity,
            "code": code,
            "field": _text(flag.get("field")) or code.removeprefix("missing_"),
            "message": _text(flag.get("text") or flag.get("label")) or "该行需要人工复核",
            "can_skip": False,
            "skip_key": f"{row_id}:{code}",
            "skip_label": "",
        })
    if record.get("_duplicate_existing"):
        identity = _joined(
            record.get("contract_no"),
            record.get("po_no"),
            record.get("customer_po"),
            record.get("product_no"),
            record.get("item_no"),
        ) or "当前订单"
        issues.append({
            "severity": "blocked",
            "code": "duplicate_existing_order",
            "field": "po_no",
            "message": f"{identity} 已存在当前客户排期；测试阶段可人工确认后重复导入",
            "can_skip": True,
            "skip_key": f"{row_id}:duplicate_existing_order",
            "skip_label": "测试阶段确认重复导入当前排期已有订单",
        })
    return issues


def _preview_row(
    *, customer_code: str, spec: HuadengCustomerMappingSpec, received_date: str,
    record: dict[str, Any], index: int, default_source_file: str, sheet_name: str,
) -> dict[str, Any]:
    values = _record_fields(customer_code, record)
    row_id = f"{customer_code}-{index}"
    issues = _issues(record, row_id)
    status = (
        "blocked" if any(item["severity"] == "blocked" for item in issues)
        else "warning" if issues else "valid"
    )
    source_file = _text(record.get("_source_po_file_name") or record.get("source_file")) or default_source_file
    source_sheet = _text(record.get("source_sheet")) or sheet_name
    source_row = _text(record.get("source_row") or record.get("row_number") or record.get("row_no"))
    source_reference = f"{source_file} · {source_sheet}"
    if source_row:
        source_reference += f" 第 {source_row} 行"
    common: dict[str, Any] = {
        "id": row_id,
        "status": status,
        "status_label": {"valid": "可导出", "warning": "需复核", "blocked": "已阻断"}[status],
        "row_role": "detail",
        "parent_product_no": "",
        "received_date": received_date,
        "po_no": "", "contract_no": "", "customer_country": "",
        "customer_name": "", "country": "", "product_no": "",
        "product_name_zh": "", "product_name_en": "", "quantity": "",
        "units_per_carton": "", "carton_count": "", "standard": "",
        "unit_price_hkd": "", "amount_hkd": "", "packaging": "",
        "line_q": "", "customer_q": "", "requested_ship_date": "",
        "input_template": spec.input_template,
        "target_template": spec.target_template,
        "item_sheet_name": sheet_name,
        "source_po_file_name": source_file,
        "lineage": {"source": source_reference, "mapping_rule": spec.rule_summary},
        "issues": issues,
    }
    for field_name, value in values.items():
        if field_name in common:
            common[field_name] = _text(value)
    if not common["customer_country"]:
        common["customer_country"] = _joined(common["customer_name"], common["country"])
    return common


def create_huadeng_customer_preview(
    *, customer_code: str, factory_id: str, received_date: str,
    po_files: list[tuple[str, bytes]], schedule_file_name: str, schedule_content: bytes,
) -> dict[str, Any]:
    spec = get_huadeng_customer_mapping(customer_code)
    if factory_id != "huadeng":
        raise HuadengCustomerOrderError(f"{spec.name} 只属于华登厂区，不能导入其他厂区")
    if not po_files:
        raise HuadengCustomerOrderError(f"请至少上传一份 {spec.name} PO")
    try:
        normalized_received_date = date.fromisoformat(received_date).isoformat()
    except ValueError as exc:
        raise HuadengCustomerOrderError("来单日期必须是 YYYY-MM-DD") from exc
    prepared = _prepare_batch(
        customer_code, po_files, schedule_file_name, schedule_content, normalized_received_date,
    )
    rows = [
        _preview_row(
            customer_code=customer_code, spec=spec, received_date=normalized_received_date,
            record=record, index=index, default_source_file=po_files[0][0],
            sheet_name=prepared.sheet_name,
        )
        for index, record in enumerate(prepared.records, start=1)
    ]
    file_names = [name for name, _ in po_files]
    po_hashes = [sha256(content).hexdigest() for _, content in po_files]
    export_position_rule = (
        "Casdon 新单按货号插入该货号最后一条历史明细之后、小计之前；"
        "外箱按同货号唯一值继承，缺失或多值时需人工填写/确认；"
        "总箱使用数量除以外箱的行公式。"
        if customer_code == "casdon"
        else "仅在对应目标表明细末尾/合计行之前插入本批新单，并从插入点向上选择最近的"
        "正常明细行继承格式、字体和公式逻辑，跳过合计/小计、分组标题和空白分隔行。"
    )
    return {
        "preview_schema_version": PREVIEW_SCHEMA_VERSION,
        "customer_code": customer_code,
        "factory_id": factory_id,
        "po_file_name": file_names[0] if len(file_names) == 1 else f"{len(file_names)}个文件",
        "po_file_names": file_names,
        "po_file_count": len(file_names),
        "schedule_file_name": schedule_file_name,
        "source_po_sha256": _combined_hash(file_names, po_hashes),
        "source_po_sha256s": po_hashes,
        "source_schedule_sha256": sha256(schedule_content).hexdigest(),
        "input_template": spec.input_template,
        "target_template": spec.target_template,
        "output_file_name": _safe_output_name(schedule_file_name, spec.name),
        "summary": {
            "total": len(rows),
            "valid": sum(row["status"] == "valid" for row in rows),
            "warning": sum(row["status"] == "warning" for row in rows),
            "blocked": sum(row["status"] == "blocked" for row in rows),
        },
        "rows": rows,
        "warnings": list(dict.fromkeys([
            spec.rule_summary,
            "导出结果完整保留当前排期的所有 Sheet、历史数据、格式、公式、图片和打印设置；"
            f"{export_position_rule}"
            "另存为新文件，不覆盖原排期。",
            *prepared.warnings,
        ])),
    }


def _export_prepared(
    *, customer_code: str, prepared: PreparedBatch, schedule_file_name: str,
    schedule_content: bytes, output_path: Path,
) -> None:
    if customer_code in {"casdon", "spin"}:
        schedule_module = casdon_schedule if customer_code == "casdon" else spin_schedule
        qty_col = schedule_module.COL["qty"]
        outer_col = schedule_module.COL["outer"]
        total_col = schedule_module.COL["total_box"]

        def row_values(record: dict[int, Any], row_no: int) -> dict[int, Any]:
            values = dict(record)
            date_columns = {
                column
                for key, column in schedule_module.COL.items()
                if key.endswith("_date") or key == "po_date"
            }
            for column in date_columns:
                value = values.get(column)
                if isinstance(value, (int, float)):
                    values[column] = datetime(1899, 12, 30) + timedelta(days=float(value))
            if values.get(qty_col) not in (None, ""):
                qty = get_column_letter(qty_col)
                outer = get_column_letter(outer_col)
                values[total_col] = f'=IF({outer}{row_no}=0,"",{qty}{row_no}/{outer}{row_no})'
            return values

        item_column = schedule_module.COL["item"]
        item_key_factory = casdon_schedule.item_base if customer_code == "casdon" else spin_schedule.item_key
        export_kwargs = {
            "filename": schedule_file_name,
            "sheet_names": (prepared.sheet_name,),
            "header_row": prepared.header_row,
            "max_col": schedule_module.EXPORT_END_COL,
            "detail_columns": (
                schedule_module.COL["contract"],
                item_column,
                schedule_module.COL["qty"],
            ),
            "row_values_factory": row_values,
        }
        insert_column_records_after_matching_groups(
            schedule_content,
            output_path,
            prepared.legacy_rows,
            existing_row_key_factory=lambda worksheet, row_no: item_key_factory(
                worksheet.cell(row_no, item_column).value
            ),
            record_key_factory=lambda record: item_key_factory(record.get(item_column)),
            **export_kwargs,
        )
        return
    if customer_code == "jakks":
        rows = jakks_new_order_writer.records_from_order(prepared.export_payload)
        append_records_to_workbook(
            schedule_content,
            output_path,
            rows,
            jakks_new_order_writer.ALIASES,
            filename=schedule_file_name,
            sheet_names=(prepared.sheet_name, "Sheet1"),
            record_overrides_formula_fields=(
                "order_date", "contact", "customer_po", "confirmation_no", "contract_no",
                "customer", "version", "item_no", "product_name", "po_description",
                "product_name_source", "quantity", "unit", "inner_pack", "outer_pack",
                "cartons", "special_notes", "ship_date", "unit_price_hkd",
                "unit_price_usd", "total_hkd", "total_usd", "country", "source_file",
            ),
        )
        return
    if customer_code == "simba":
        aliases = {
            field: tuple(dict.fromkeys(
                [title] + [alias for alias, target in simba_schedule.FIELD_ALIASES.items() if target == field]
            ))
            for field, title in simba_schedule.FIELD_TITLES.items()
        }

        def item_total_values(
            _records: list[dict[str, Any]],
            start_row: int,
            end_row: int,
            _total_row: int,
            column_map: dict[int, str],
        ) -> dict[int, Any]:
            field_columns = {field: column for column, field in column_map.items()}
            name_column = field_columns["product_name"]
            quantity_column = field_columns["quantity"]
            quantity_letter = get_column_letter(quantity_column)
            return {
                name_column: "合计：",
                quantity_column: (
                    f"=SUM({quantity_letter}{start_row}:{quantity_letter}{end_row})"
                ),
            }

        append_records_to_workbook(
            schedule_content,
            output_path,
            prepared.records,
            aliases,
            filename=schedule_file_name,
            sheet_names=(prepared.sheet_name,),
            group_total_key_factory=lambda record: _text(record.get("item_no")).upper(),
            group_total_row_values_factory=item_total_values,
        )
        return
    append_records_to_workbook(
        schedule_content, output_path, prepared.records,
        spin_master_new_order_writer.FIELD_ALIASES,
        filename=schedule_file_name, sheet_names=(prepared.sheet_name, "SPIN总汇"),
    )


def export_huadeng_customer_schedule(
    *, customer_code: str, factory_id: str, received_date: str,
    po_files: list[tuple[str, bytes]], schedule_file_name: str, schedule_content: bytes,
    skipped_issue_keys: set[str] | None = None,
    manual_overrides: list[dict[str, str]] | None = None,
) -> tuple[bytes, str, dict[str, Any]]:
    preview = create_huadeng_customer_preview(
        customer_code=customer_code, factory_id=factory_id, received_date=received_date,
        po_files=po_files, schedule_file_name=schedule_file_name,
        schedule_content=schedule_content,
    )
    decorate_manual_resolution_policy(preview)
    apply_overrides_to_preview(preview, manual_overrides or [])
    requested_skips = set(skipped_issue_keys or set())
    skippable = {
        issue["skip_key"]
        for row in preview["rows"]
        for issue in row["issues"]
        if issue.get("skip_key")
    }
    unknown_skips = requested_skips - skippable
    if unknown_skips:
        raise HuadengCustomerOrderError("确认项已失效，请重新预览后再生成")
    blockers = [
        issue["message"] for row in preview["rows"] for issue in row["issues"]
        if issue["severity"] == "blocked" and issue["skip_key"] not in requested_skips
    ]
    if blockers:
        raise HuadengCustomerOrderError("仍有阻断项：" + "；".join(dict.fromkeys(blockers)))
    if not preview["rows"]:
        raise HuadengCustomerOrderError("本批文件没有可安全生成的新单明细，请查看预览告警")
    prepared = _prepare_batch(
        customer_code, po_files, schedule_file_name, schedule_content, received_date,
    )
    field_aliases = {
        "jakks": {
            "product_name_zh": "product_name",
            "product_name_en": "po_description",
            "units_per_carton": "outer_pack",
            "carton_count": "cartons",
            "packaging": "special_note",
            "line_q": "inspection_date",
            "requested_ship_date": "ship_date",
        },
        "simba": {
            "product_name_zh": "product_name",
            "product_name_en": "english_name",
            "units_per_carton": "outer_pack",
            "carton_count": "cartons",
            "packaging": "special_remark",
            "line_q": "inspection_date",
            "requested_ship_date": "po_ship_date",
        },
        "spin-master": {
            "product_name_zh": "product_name",
            "product_name_en": "english_name",
            "units_per_carton": "outer_pack",
            "carton_count": "cartons",
            "packaging": "special_notes",
            "line_q": "inspection_date",
            "requested_ship_date": "ship_date",
        },
    }.get(customer_code, {})
    apply_overrides_to_records(
        prepared.records,
        [str(row["id"]) for row in preview["rows"]],
        manual_overrides or [],
        field_aliases=field_aliases,
    )
    if customer_code in {"casdon", "spin"}:
        schedule_module = casdon_schedule if customer_code == "casdon" else spin_schedule
        column_by_field = {
            "po_no": "customer_po",
            "contract_no": "contract",
            "customer_name": "customer",
            "country": "ship_country",
            "product_no": "item",
            "product_name_zh": "cn_name",
            "product_name_en": "english_name",
            "quantity": "qty",
            "units_per_carton": "outer",
            "standard": "version",
            "unit_price_hkd": "unit_hkd",
            "amount_hkd": "total_hkd",
            "packaging": "special_note",
            "line_q": "inspection_date",
            "requested_ship_date": "ship_date",
        }
        for record in prepared.records:
            values = record.get("_legacy_values")
            if not isinstance(values, dict):
                continue
            for field, column_key in column_by_field.items():
                if field in record:
                    values[schedule_module.COL[column_key]] = record[field]
    with TemporaryDirectory(prefix=f"huadeng-{customer_code}-output-") as temp_dir:
        output_path = Path(temp_dir) / preview["output_file_name"]
        try:
            _export_prepared(
                customer_code=customer_code, prepared=prepared,
                schedule_file_name=schedule_file_name, schedule_content=schedule_content,
                output_path=output_path,
            )
        except Exception as exc:
            name = get_huadeng_customer_mapping(customer_code).name
            raise HuadengCustomerOrderError(f"生成 {name} 新单失败：{exc}") from exc
        return output_path.read_bytes(), preview["output_file_name"], preview
