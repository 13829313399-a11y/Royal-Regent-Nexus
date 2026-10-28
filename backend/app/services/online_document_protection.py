"""Bounded, opt-in XLSX input ranges. Preserve unrelated OOXML package parts.

Sheet passwords are an editor affordance, not the authorization boundary: every
saved workbook is compared with an immutable protected baseline, for all actors.
"""
from copy import deepcopy
from decimal import Decimal, InvalidOperation
from io import BytesIO
import posixpath
import re
from uuid import uuid4
from zipfile import ZipFile

from fastapi import HTTPException
from lxml import etree as ET
from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries, get_column_letter
from openpyxl.utils.protection import hash_password

from app.services import online_document_protocol as protocol

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
Q = lambda name: f"{{{NS}}}{name}"
MAX_INPUTS = 50000
MAX_CELLS = 500000


def fail(message="此表格包含暂不支持的保护结构，请保持原文件，另存为普通 XLSX 后重试"):
    raise HTTPException(422, message)


def xml(data):
    parser = ET.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False)
    try:
        root = ET.fromstring(data, parser)
        if root.getroottree().docinfo.doctype:
            fail()
        return root
    except ET.XMLSyntaxError:
        fail()


def semantic(node):
    if node is None:
        return None
    return (node.tag, tuple(sorted(node.attrib.items())), node.text or "", tuple(semantic(c) for c in node))


class Book:
    def __init__(self, data):
        protocol.validate(data, "xlsx")
        self.archive = ZipFile(BytesIO(data))
        names = self.archive.namelist()
        if any(n.startswith("_xmlsignatures/") for n in names):
            fail("带数字签名的文件不能直接设置保护，请先另存未签名副本")
        self.workbook = xml(self.archive.read("xl/workbook.xml"))
        if self.workbook.tag != Q("workbook"):
            fail()
        rels = xml(self.archive.read("xl/_rels/workbook.xml.rels"))
        self.external_targets = tuple(sorted((r.get("Id"), r.get("Target"), r.get("TargetMode", "Internal"))
                                             for r in rels if r.get("Type", "").endswith("/externalLink")))
        targets = {}
        for r in rels:
            if r.get("Type", "").endswith("/worksheet"):
                target = r.get("Target", "")
                target = posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)
                if r.get("TargetMode") == "External" or not target.startswith("xl/"):
                    fail()
                targets[r.get("Id")] = target
        sheets = self.workbook.find(Q("sheets"))
        if sheets is None or not len(sheets) or len(sheets) > 100:
            fail()
        self.sheets = []
        self.cells = {}
        seen = set()
        strings = []
        if "xl/sharedStrings.xml" in names:
            strings = ["".join(si.itertext(Q("t"))) for si in xml(self.archive.read("xl/sharedStrings.xml"))]
        self.styles = xml(self.archive.read("xl/styles.xml"))
        self.xfs = self.styles.find(Q("cellXfs"))
        if self.xfs is None or not len(self.xfs) or len(self.xfs) > 64000:
            fail()
        self.values = {}
        self.formulas = {}
        for s in sheets:
            name, path = s.get("name"), targets.get(s.get(f"{{{REL}}}id"))
            if not name or name in seen or not path:
                fail()
            seen.add(name)
            root = xml(self.archive.read(path))
            if root.tag != Q("worksheet"):
                fail()
            self.sheets.append((name, path, root, s.get("state", "visible")))
            self.cells[name] = {}
            formulas, shared, pending = {}, {}, []
            sheet_data = root.find(Q("sheetData"))
            if sheet_data is None:
                fail()
            seen_rows = set()
            for row in sheet_data:
                row_id = int(row.get("r", "0"))
                if row_id in seen_rows or not 1 <= row_id <= 1048576:
                    fail()
                seen_rows.add(row_id)
                for cell in row.findall(Q("c")):
                    address = cell.get("r", "")
                    if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", address):
                        fail()
                    cr, cc = coordinate_to_tuple(address)
                    if cr != row_id or cc > 16384 or address in self.cells[name]:
                        fail()
                    self.cells[name][address] = cell
                    f = cell.find(Q("f"))
                    v = cell.find(Q("v"))
                    if f is not None:
                        kind = f.get("t", "normal")
                        if kind not in {"normal", "shared"}:
                            fail("数组公式、动态数组或数据表公式暂不支持区域保护；普通编辑不受影响")
                        if kind == "shared":
                            if f.text:
                                if f.get("si") in shared:
                                    fail()
                                shared[f.get("si")] = (address, f.text)
                            pending.append((address, f.get("si")))
                        else:
                            formulas[address] = f.text or ""
                    else:
                        t = cell.get("t", "n")
                        value = v.text if v is not None else None
                        if t == "s":
                            value, t = strings[int(value)], "s"
                        elif t == "inlineStr":
                            inline = cell.find(Q("is"))
                            value, t = "".join(inline.itertext(Q("t"))) if inline is not None else "", "s"
                        elif t in {"n", "b"} and value not in {None, ""}:
                            value = Decimal(value)
                            if not value.is_finite():
                                fail()
                        if value not in {None, ""}:
                            self.values[(name, address)] = (t, value)
            for address, index in pending:
                if index not in shared:
                    fail()
                origin, formula = shared[index]
                formulas[address] = Translator("=" + formula, origin=origin).translate_formula(address)[1:]
            self.formulas.update({(name, a): f for a, f in formulas.items()})
            if sum(len(v) for v in self.cells.values()) > MAX_CELLS:
                fail("区域保护最多检查 50 万个已使用单元格，请先拆分大表")
        # Names and table/external reference definitions can change a formula's
        # meaning even when its literal text is unchanged.
        defined_names = self.workbook.find(Q("definedNames"))
        # ONLYOFFICE omits an empty definedNames container on save. Its absence
        # and an attribute-free, child-free container define the same names.
        if defined_names is not None and not len(defined_names) and not defined_names.attrib and not (defined_names.text or "").strip():
            defined_names = None
        self.references = (semantic(defined_names), {
            n: semantic(xml(self.archive.read(n))) for n in names
            if n.startswith(("xl/tables/", "xl/externalLinks/")) and n.endswith((".xml", ".rels"))})
        workbook_pr = self.workbook.find(Q("workbookPr"))
        calc = self.workbook.find(Q("calcPr"))
        self.calculation = (
            (workbook_pr.get("date1904", "0") if workbook_pr is not None else "0") in {"1", "true"},
            tuple((key, calc.get(key, default) if calc is not None else default) for key, default in (
                ("calcMode", "auto"), ("refMode", "A1"), ("fullPrecision", "1"),
                ("iterate", "0"), ("iterateCount", "100"), ("iterateDelta", "0.001"))),
            semantic(self.workbook.find(Q("externalReferences"))))
        self.structure = {}
        for name, path, root, _ in self.sheets:
            relpath = posixpath.dirname(path) + "/_rels/" + posixpath.basename(path) + ".rels"
            table_targets = []
            if relpath in names:
                for r in xml(self.archive.read(relpath)):
                    if r.get("Type", "").endswith("/table"):
                        table_targets.append((r.get("Id"), r.get("Target"), r.get("TargetMode", "Internal")))
            self.structure[name] = (
                tuple(semantic(root.find(Q(tag))) for tag in ("mergeCells", "autoFilter", "tableParts", "dataValidations")),
                tuple((r.get("r"), r.get("hidden")) for r in root.find(Q("sheetData")) if r.get("hidden", "0") in {"1", "true"}),
                tuple(sorted(table_targets)))

    def locked(self, sheet, address):
        cell = self.cells[sheet].get(address)
        index = int(cell.get("s", "0")) if cell is not None else 0
        if not 0 <= index < len(self.xfs):
            fail()
        protection = self.xfs[index].find(Q("protection"))
        return protection is None or protection.get("locked", "1") not in {"0", "false"}

    def patch(self, parts):
        result = BytesIO()
        with ZipFile(result, "w") as target:
            for entry in self.archive.infolist():
                target.writestr(entry, ET.tostring(parts[entry.filename], encoding="UTF-8", xml_declaration=True)
                                if entry.filename in parts else self.archive.read(entry.filename))
        return result.getvalue()


def read(data):
    try:
        return Book(data)
    except (ValueError, KeyError, IndexError, TypeError, InvalidOperation, ET.LxmlError, TranslatorError):
        fail()


def ranges_for(book, rules):
    if len(rules) != len(book.sheets) or {r["sheet"] for r in rules} != {s[0] for s in book.sheets}:
        fail("请为每个工作表设置填写区域；留空表示整张工作表只读")
    result, count = {}, 0
    for rule in rules:
        addresses = set()
        for value in rule["ranges"]:
            if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}(:[A-Z]{1,3}[1-9][0-9]{0,6})?", value):
                fail("填写区域格式应为 B2:D100，多个区域用英文逗号分隔")
            c1, r1, c2, r2 = range_boundaries(value)
            if not (1 <= c1 <= c2 <= 16384 and 1 <= r1 <= r2 <= 1048576):
                fail("填写区域超出工作表范围")
            count += (r2-r1+1)*(c2-c1+1)
            if count > MAX_INPUTS:
                fail("填写区域合计最多 50000 个单元格，请缩小范围")
            addresses.update(f"{get_column_letter(c)}{r}" for r in range(r1, r2+1) for c in range(c1, c2+1))
        result[rule["sheet"]] = addresses
    return result


def inspect(data):
    book = read(data)
    return [{"name": name, "formula_count": sum(s == name for s, _ in book.formulas),
             "used_cells": len(book.cells[name])} for name, _, _, _ in book.sheets]


def enable(data, rules):
    book = read(data)
    allowed = ranges_for(book, rules)
    if sum(len(set(book.cells[name]) | allowed[name]) for name in book.cells) > MAX_CELLS:
        fail("设置填写区域后将超过 50 万个单元格，请缩小范围或拆分大表")
    original = book.workbook.find(Q("workbookProtection"))
    if (original is not None and any(v not in {"0", "false"} for v in original.attrib.values())) or any(
            root.find(Q("sheetProtection")) is not None or root.find(Q("protectedRanges")) is not None for _, _, root, _ in book.sheets):
        fail("文件已有工作簿或工作表保护，请由拥有者解除原保护并保存后再设置")
    if original is not None:
        book.workbook.remove(original)
    password = hash_password(uuid4().hex)
    parts = {"xl/styles.xml": book.styles, "xl/workbook.xml": book.workbook}
    # Normalize existing styles to locked, then clone only styles actually used
    # in input cells. Original XML parts such as drawings remain byte-for-byte.
    for xf in book.xfs:
        p = xf.find(Q("protection"))
        if p is None:
            p = ET.SubElement(xf, Q("protection"))
        p.set("locked", "1"); xf.set("applyProtection", "1")
    unlocked_styles = {}
    for name, path, root, _ in book.sheets:
        sheet_data = root.find(Q("sheetData"))
        rows = {int(r.get("r")): r for r in sheet_data}
        for address in sorted(allowed[name], key=coordinate_to_tuple):
            if (name, address) in book.formulas:
                continue
            r, _ = coordinate_to_tuple(address)
            if r not in rows:
                rows[r] = ET.SubElement(sheet_data, Q("row"), r=str(r))
            cell = book.cells[name].get(address)
            if cell is None:
                cell = ET.SubElement(rows[r], Q("c"), r=address)
            style = int(cell.get("s", "0"))
            if not 0 <= style < len(book.xfs):
                fail()
            if style not in unlocked_styles:
                clone = deepcopy(book.xfs[style]); clone.find(Q("protection")).set("locked", "0")
                unlocked_styles[style] = len(book.xfs); book.xfs.append(clone)
            cell.set("s", str(unlocked_styles[style]))
        for row in rows.values():
            row[:] = sorted(row, key=lambda c: coordinate_to_tuple(c.get("r")))
        sheet_data[:] = [rows[r] for r in sorted(rows)]
        coordinates = [coordinate_to_tuple(c.get("r")) for row in sheet_data for c in row]
        if coordinates:
            dimension = root.find(Q("dimension"))
            if dimension is not None:
                dimension.set("ref", f"{get_column_letter(min(c for r, c in coordinates))}{min(r for r, c in coordinates)}:{get_column_letter(max(c for r, c in coordinates))}{max(r for r, c in coordinates)}")
        p = ET.Element(Q("sheetProtection"), sheet="1", password=password, objects="1", scenarios="1",
                       formatCells="1", formatColumns="1", formatRows="1", insertColumns="1", insertRows="1",
                       insertHyperlinks="1", deleteColumns="1", deleteRows="1", sort="1", autoFilter="1", pivotTables="1",
                       selectLockedCells="0", selectUnlockedCells="0")
        # OOXML order: sheetCalcPr precedes sheetProtection.
        previous = root.find(Q("sheetCalcPr"))
        root.insert(root.index(previous if previous is not None else sheet_data)+1, p)
        parts[path] = root
    book.xfs.set("count", str(len(book.xfs)))
    if len(book.xfs) > 64000:
        fail("文件样式过多，请简化格式后设置保护")
    protection = ET.Element(Q("workbookProtection"), lockStructure="1", workbookPassword=password)
    sheets = book.workbook.find(Q("sheets"))
    views = book.workbook.find(Q("bookViews"))
    book.workbook.insert(book.workbook.index(views if views is not None else sheets), protection)
    return book.patch(parts)


def disable(data):
    book = read(data)
    parts = {"xl/workbook.xml": book.workbook}
    p = book.workbook.find(Q("workbookProtection"))
    if p is not None:
        book.workbook.remove(p)
    for _, path, root, _ in book.sheets:
        p = root.find(Q("sheetProtection"))
        if p is not None:
            root.remove(p)
        parts[path] = root
    return book.patch(parts)


def validate_saved(baseline, candidate, rules):
    before, after = read(baseline), read(candidate)
    allowed = ranges_for(before, rules)
    def reject():
        fail("保存被阻止：公式、锁定区域或表格结构发生变化。请撤销相关修改后重试；需要调整模板时请先关闭所有编辑窗口，由拥有者解除保护")
    if [(n, s) for n, _, _, s in before.sheets] != [(n, s) for n, _, _, s in after.sheets]:
        reject()
    if (before.formulas != after.formulas or before.references != after.references or
            before.calculation != after.calculation or before.structure != after.structure or before.external_targets != after.external_targets):
        reject()
    if semantic(before.workbook.find(Q("workbookProtection"))) != semantic(after.workbook.find(Q("workbookProtection"))):
        reject()
    for old, new in zip(before.sheets, after.sheets):
        name = old[0]
        if semantic(old[2].find(Q("sheetProtection"))) != semantic(new[2].find(Q("sheetProtection"))):
            reject()
        if new[2].find(Q("protectedRanges")) is not None:
            reject()
        for address in set(before.cells[name]) | set(after.cells[name]):
            editable = address in allowed[name] and (name, address) not in before.formulas
            if not editable:
                if before.values.get((name, address)) != after.values.get((name, address)):
                    reject()
                if address in after.cells[name] and not after.locked(name, address):
                    reject()
        # The immutable whitelist remains the authority even if a client changes
        # style defaults or adds empty cells; every later save is checked again.
