"""Opening stock and physical counts, on the existing append-only warehouse ledger.

Metadata events never affect stock. OPENING and STOCKTAKE children carry each
physical quantity once. Purchase receipt progress and supplier reconciliation do
not include either kind.
"""
import copy
import io
import json
import re
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import PurePath
from uuid import UUID, uuid4, uuid5

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select

from app.core.time import business_now
from app.models.warehouse_operations import WarehouseOperation
from app.services import warehouse_operations as ops
from app.services.fabric_procurement_parser import digest, dump
from app.services.fabric_receiving import quantity

META = {'OPENING_IMPORT', 'STOCKTAKE_DRAFT', 'STOCKTAKE_SAVE', 'STOCKTAKE_REFRESH', 'STOCKTAKE_POST', 'STOCKTAKE_CANCEL'}
HEADERS = ['来源行号', '物料编码', '物料名称', '物料分类', '单位', '仓库', '仓位', '缸号', '卷号', '期初数量', '来源单位', '备注', '货号', '客户名称', '物料明细']
KEYS = ['source_line', 'item_code', 'item_name', 'material_category', 'unit', 'warehouse', 'location_code', 'lot', 'roll_no', 'quantity', 'counterparty', 'reason', 'style_no', 'customer_name', 'material_detail']
CATEGORIES = {'布料': 'FABRIC', '辅料': 'ACCESSORY', '线': 'THREAD'}
MAX_BYTES = 2 * 1024 * 1024


def template(db):
    book = Workbook()
    sheet = book.active
    sheet.title = '期初库存'
    sheet.append(HEADERS)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = 'A1:O1001'
    for cell in sheet[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='0F766E')
    for index, width in enumerate([16, 24, 32, 14, 12, 22, 20, 20, 20, 18, 28, 40, 24, 28, 48], 1):
        from openpyxl.utils import get_column_letter
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2, max_row=1001):
        for cell in row:
            cell.number_format = '@'
    guide = book.create_sheet('填写说明')
    for line in ['每行一项实物批次，最多1000行。请保留表头；不要插入公式。',
                 '仅导入尚未在本系统登记的历史结余，已收料的实物不能再次导入。',
                 '来源清单编号、库存截止日由页面统一填写；同一清单同一来源行号不能重复。',
                 '编码、缸号和卷号按文本填写，保留前导零。分类填写：布料、辅料、线。',
                 '货号、客户名称、物料明细必填；货号填写产品/款号，物料明细填写规格、颜色、成分等实物说明。',
                 '客户和货号记录本行库存的来源/用途，不改变共用物料资料；共用库存须填写经核实的共用说明。',
                 '数量为正数，最多六位小数；不同单位不自动换算。布料必须填写缸号。',
                 '仓库与仓位填写基础资料中的名称/编号。须与物料分类匹配，停用或未分类仓位不能导入。',
                 '预览不会入账；确认后建立期初批次，不增加采购实收及供应商本期进货。',
                 '相同来源及内容重复导入会跳过；有冲突的行须修正后整批重新预览。']:
        guide.append([line])
    guide.column_dimensions['A'].width = 110
    locations = book.create_sheet('可用仓位')
    locations.append(['仓库', '仓位'])
    for item in ops.locations.catalog(db, 'fabric'):
        if item['status'] == 'ACTIVE':
            locations.append([item['warehouse'], item['code']])
            for cell in locations[locations.max_row]:
                cell.data_type = 's'
                cell.number_format = '@'
    locations.column_dimensions['A'].width = 30
    locations.column_dimensions['B'].width = 30
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def parse(filename, content):
    if not filename.lower().endswith('.xlsx') or not content or len(content) > MAX_BYTES:
        raise HTTPException(422, '请选择不超过 2 MiB 的 .xlsx 期初库存模板')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(part.file_size for part in archive.infolist()) > 20 * MAX_BYTES or len(archive.infolist()) > 200:
                raise ValueError()
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception:
        raise HTTPException(422, '无法读取模板，文件损坏或解压后过大')
    try:
        if '期初库存' not in book.sheetnames:
            raise HTTPException(422, '缺少“期初库存”工作表，请使用下载的模板')
        sheet = book['期初库存']
        if (sheet.max_row or 0) > 1001 or (sheet.max_column or 0) > len(HEADERS):
            raise HTTPException(422, '模板超出1000行或标准列范围，请删除范围外的内容和格式后分批导入')
        rows = sheet.iter_rows(max_row=1002, max_col=len(HEADERS) + 1)
        header = next(rows)
        if [cell.value for cell in header[:len(HEADERS)]] != HEADERS or header[-1].value is not None:
            raise HTTPException(422, '模板表头不匹配，请保留下载模板的列名和顺序')
        result = []
        for row_number, cells in enumerate(rows, 2):
            if all(cell.value in (None, '') for cell in cells):
                continue
            if row_number > 1001:
                raise HTTPException(422, '每次最多导入1000行，请分批整理')
            errors = []
            if cells[-1].value is not None:
                errors.append('模板列以外存在数据')
            values = []
            for index, cell in enumerate(cells[:len(HEADERS)]):
                value = cell.value
                if cell.data_type == 'f':
                    errors.append(f'{HEADERS[index]}不能使用公式')
                if index in (1, 7, 8, 12) and value is not None and not isinstance(value, str):
                    errors.append(f'{HEADERS[index]}请设为文本后填写，防止丢失前导零')
                value = '' if value is None else str(value).strip()
                if len(value) > (1000 if index in (11, 14) else 255):
                    errors.append(f'{HEADERS[index]}过长')
                values.append(value)
            result.append(dict(zip(KEYS, values), row_number=row_number, errors=errors))
        if not result:
            raise HTTPException(422, '模板中没有待导入的实物记录')
        return result
    finally:
        book.close()


def physical_key(row):
    return tuple(row.get(key, '') for key in ('item_code', 'unit', 'location_id', 'lot', 'roll_no'))


def inventory_token(stock, entries):
    # Include event identities as well as quantities: issue+reversal must invalidate
    # a count even when the numerical balance returns to its prior value.
    return digest([stock, [(e['id'], e['sequence']) for e in entries.values() if e['kind'] not in META]])


def count_tokens(stock, entries, batch_ids):
    """Bind only selected physical batches and events affecting their evidence."""
    selected = set(batch_ids)
    events = {identifier: [] for identifier in selected}
    receipt_batches = {}
    for identifier in selected:
        receipt = (stock.get(identifier, {}).get('receipt_origin') or {}).get('id')
        if receipt:
            receipt_batches.setdefault(receipt, set()).add(identifier)
    affected_by_entry = {}
    for entry in entries.values():
        body, kind = entry['data'], entry['kind']
        if kind in META:
            continue
        affected = {entry['id'], body.get('batch_id')}
        if kind in {'RETURN', 'CORRECTION'}:
            affected.update(affected_by_entry.get(body.get('original_id'), set()))
        if kind == 'CORRECTION':
            affected.update(receipt_batches.get(body.get('procurement_original_receipt_id'), set()))
        affected_by_entry[entry['id']] = affected
        for identifier in selected.intersection(affected):
            events[identifier].append((entry['id'], entry['sequence']))
    return {identifier: digest([stock.get(identifier), events[identifier]]) for identifier in sorted(selected)}


def set_count_scope(document, stock, entries):
    document['count_scope'] = 'BATCHES_V1'
    document['batch_tokens'] = count_tokens(stock, entries, [line['batch_id'] for line in document['lines']])


def count_changes(document, stock, entries):
    ids = [line['batch_id'] for line in document['lines']]
    if document.get('count_scope') != 'BATCHES_V1':
        # Old drafts retain their original conservative contract until explicitly refreshed.
        return ids if document['ledger_token'] != inventory_token(stock, entries) else []
    current = count_tokens(stock, entries, ids)
    return [identifier for identifier in ids if current[identifier] != document['batch_tokens'].get(identifier)]


def check_date(value):
    if value > business_now().date():
        raise HTTPException(422, '库存日期不能晚于今天')


def opening_preview(db, actor, filename, content, cutoff, source_no):
    ops.ensure_schema(db)
    check_date(cutoff)
    if not source_no.strip() or len(source_no.strip()) > 128:
        raise HTTPException(422, '请填写不超过128字的原始库存清单编号')
    rows = parse(filename, content)
    stock, entries = ops.projection(db, 'fabric')
    places = ops.locations.catalog(db, 'fabric')
    from app.services.fabric_master import records
    masters = records(db)
    used, physical = set(), set()
    normalized = []
    for raw in rows:
        row = {**raw, 'source_no': source_no.strip(), 'business_date': cutoff.isoformat(), 'source_system': 'OPENING', 'process_state': ''}
        row['material_category'] = CATEGORIES.get(row['material_category'], row['material_category'])
        if row['material_category'] not in CATEGORIES.values():
            row['errors'].append('物料分类须为布料、辅料或线')
        for key, label in [('source_line', '来源行号'), ('item_code', '物料编码'), ('item_name', '物料名称'), ('unit', '单位'),
                           ('style_no', '货号'), ('customer_name', '客户名称'), ('material_detail', '物料明细')]:
            if not row[key]:
                row['errors'].append(f'请填写{label}')
        if any(len(row[key]) > limit for key, limit in [('source_line',128),('item_code',128),('unit',32),('lot',128),('roll_no',128)]):
            row['errors'].append('编码、单位或批次标识过长')
        if row['material_category'] == 'FABRIC' and not row['lot']:
            row['errors'].append('布料必须填写缸号')
        if row['material_category'] != 'FABRIC' and (row['lot'] or row['roll_no']):
            row['errors'].append('辅料和线不填写布料缸号或卷号')
        place = next((p for p in places if p['warehouse'] == row['warehouse'] and p['code'] == row['location_code'] and p['status'] == 'ACTIVE'), None)
        row.update(location_id=place['id'] if place else '', location=place['label'] if place else '', location_snapshot=place or {})
        if not place:
            row['errors'].append('仓库或仓位不存在/已停用，请先维护基础资料')
        elif not place.get('category'):
            row['errors'].append('仓位尚未分类，请先在基础资料设置对应原料分区')
        elif place['category'] != row['material_category']:
            row['errors'].append('仓位与物料分类不匹配，请选择对应原料分区')
        try:
            row['quantity'] = ops.text(quantity(row['quantity'], '期初数量'))
        except HTTPException as exc:
            row['errors'].append(exc.detail)
        for master in masters:
            if master.kind == 'MATERIAL' and master.code == row['item_code']:
                fields = json.loads(master.data_json)
                if master.status != 'ACTIVE' or master.name != row['item_name'] or fields.get('unit') != row['unit'] or fields.get('category') != row['material_category']:
                    row['errors'].append('物料与已启用基础资料的名称、单位或分类不一致')
            if master.kind == 'UNIT' and row['unit'] in (master.code, master.name) and master.status != 'ACTIVE':
                row['errors'].append('计量单位未启用')
        key = physical_key(row)
        if row['source_line'] in used or key in physical:
            row['errors'].append('文件内来源行号或实物批次重复，请合并后导入')
        used.add(row['source_line']); physical.add(key)
        identity_fields = ('item_code','item_name','material_category','unit','location_id','lot','roll_no','quantity','counterparty','reason','style_no','customer_name','material_detail')
        row['line_fingerprint'] = digest([cutoff.isoformat(), {k: row[k] for k in identity_fields}])
        existing = next((e for e in entries.values() if e['kind'] == 'OPENING' and e['data']['source_no'] == source_no.strip() and e['data']['source_line'] == row['source_line']), None)
        if existing:
            row['action'] = 'SKIP'
            if existing['data']['line_fingerprint'] != row['line_fingerprint']:
                row['errors'].append('该来源行已入账且内容不同，请用盘点核对差异')
        else:
            row['action'] = 'NEW'
            if any(physical_key(batch) == key for batch in stock.values()):
                row['errors'].append('该实物批次已有库存记录，请用盘点核对，不能重复导入期初')
        if row['errors']:
            row['action'] = 'ERROR'
        normalized.append(row)
    token = digest(['fabric-opening-v1', actor.id, content.hex(), cutoff.isoformat(), source_no.strip(), inventory_token(stock, entries), normalized])
    totals = {}
    for row in normalized:
        if row['action'] == 'NEW':
            totals[row['unit']] = ops.text(Decimal(totals.get(row['unit'], '0')) + Decimal(row['quantity']))
    return dict(filename=PurePath(filename.replace('\\','/')).name, rows=normalized, token=token, totals=totals,
        new_count=sum(r['action']=='NEW' for r in normalized), skip_count=sum(r['action']=='SKIP' for r in normalized), error_count=sum(r['action']=='ERROR' for r in normalized))


def replay(db, actor, request_id, signature):
    old = db.scalar(select(WarehouseOperation).where(WarehouseOperation.factory_id == ops.FACTORY, WarehouseOperation.warehouse == 'fabric', WarehouseOperation.request_id == str(request_id)))
    if old:
        if old.request_hash != signature or old.actor_id != str(actor.id):
            raise HTTPException(409, '请求编号已用于其他内容或操作者')
        return json.loads(old.payload_json)['result']


def append(db, actor, stock, entries, kind, body, business_date, request_id, signature):
    identifier = 'FB-' + uuid4().hex[:20]
    revision = max((e['sequence'] for e in entries.values()), default=0) + 1
    now = business_now().isoformat()
    entry = dict(id=identifier, kind=kind, sequence=revision, business_date=str(business_date), data=body, actor_name=actor.display_name, occurred_at=now)
    ops.apply_entry(stock, entries, entry, 'fabric')
    row = WarehouseOperation(id=identifier, factory_id=ops.FACTORY, warehouse='fabric', sequence=revision, request_id=str(request_id), request_hash=signature,
        kind=kind, business_date=str(business_date), payload_json=dump(body), actor_id=str(actor.id), actor_name=actor.display_name, occurred_at=now)
    db.add(row)
    return row


def apply_opening(db, actor, filename, content, cutoff, source_no, request_id, token, confirmed):
    ops.ensure_schema(db)
    signature = digest(['fabric-opening-v1', str(actor.id), digest(content.hex()), cutoff.isoformat(), source_no, token, confirmed])
    try:
        state = ops.procurement.lock_factory(db)
        previous = replay(db, actor, request_id, signature)
        if previous is not None:
            db.rollback(); return previous
        if not confirmed:
            raise HTTPException(422, '请确认仅导入未登记的历史实物结余')
        preview = opening_preview(db, actor, filename, content, cutoff, source_no)
        if preview['token'] != token:
            raise HTTPException(409, '文件、基础资料或库存已变化，请重新预览核对')
        if preview['error_count']:
            raise HTTPException(422, '请修正全部错误后重新预览，本次未导入任何库存')
        stock, entries = ops.projection(db, 'fabric')
        header = append(db, actor, stock, entries, 'OPENING_IMPORT', {}, cutoff, request_id, signature)
        ids = []
        for index, item in enumerate(preview['rows']):
            if item['action'] != 'NEW':
                continue
            body = {k: v for k,v in item.items() if k not in {'errors','action'}}
            body.update(import_id=header.id, filename=preview['filename'], file_hash=digest(content.hex()), allocations=[])
            child = append(db, actor, stock, entries, 'OPENING', body, cutoff, uuid5(UUID(str(request_id)), f'opening-{index}'), signature)
            ids.append(child.id)
        result = dict(id=header.id, business_date=cutoff.isoformat(), source_no=source_no, filename=preview['filename'], actor_name=actor.display_name,
            posted_count=len(ids), skipped_count=preview['skip_count'], totals=preview['totals'], rows=preview['rows'], batch_ids=ids)
        header.payload_json = dump({'result':result})
        state.revision += 1
        db.commit()
        return result
    except Exception:
        db.rollback(); raise


def count_documents(entries):
    result = {}
    for entry in entries.values():
        if entry['kind'].startswith('STOCKTAKE_') and entry['kind'] in META:
            record = entry['data'].get('result')
            if record:
                result[record['id']] = record
    return result


def count_snapshot(stock, ids, entries):
    if len(set(ids)) != len(ids):
        raise HTTPException(422, '同一批次不能重复盘点')
    rows = []
    for identifier in ids:
        batch = ops.batch_for(stock, identifier)
        if entries.get(identifier, {}).get('reversed') or any(e['kind']=='CORRECTION' and e['data'].get('procurement_original_receipt_id') == batch.get('receipt_origin',{}).get('id') and e['data'].get('procurement_original_receipt_id') for e in entries.values()):
            raise HTTPException(409, '已冲销入库的批次不能通过盘点恢复，请核对来源后重新收料')
        if not batch.get('location_id'):
            raise HTTPException(409, '请先核实所选批次的正式仓位')
        rows.append(dict(batch_id=identifier, snapshot=copy.deepcopy(batch), actual='', reason=''))
    return rows


def create_count(db, actor, payload):
    ops.ensure_schema(db)
    signature = digest(['fabric-count-create', str(actor.id), payload.model_dump(mode='json')])
    try:
        state = ops.procurement.lock_factory(db)
        old = replay(db, actor, payload.request_id, signature)
        if old is not None:
            db.rollback(); return old
        stock, entries = ops.projection(db, 'fabric')
        lines = count_snapshot(stock, payload.batch_ids, entries)
        token = inventory_token(stock, entries)
        today = business_now().date().isoformat()
        row = append(db, actor, stock, entries, 'STOCKTAKE_DRAFT', {}, today, payload.request_id, signature)
        result = dict(id=row.id, revision=1, status='DRAFT', actor_name=actor.display_name, created_by=str(actor.id), business_date=today,
            reason=payload.reason, lines=lines, ledger_token=token, events=[dict(action='CREATE', actor_name=actor.display_name, at=row.occurred_at)])
        set_count_scope(result, stock, entries)
        row.payload_json = dump({'result':result})
        state.revision += 1
        db.commit(); return result
    except Exception:
        db.rollback(); raise


def act_count(db, actor, identifier, payload):
    ops.ensure_schema(db)
    signature = digest(['fabric-count-act', identifier, str(actor.id), payload.model_dump(mode='json')])
    try:
        state = ops.procurement.lock_factory(db)
        old = replay(db, actor, payload.request_id, signature)
        if old is not None:
            db.rollback(); return old
        stock, entries = ops.projection(db, 'fabric')
        document = copy.deepcopy(count_documents(entries).get(identifier))
        if not document:
            raise HTTPException(404, '盘点单不存在于本仓')
        if document['revision'] != payload.expected_revision or document['status'] != 'DRAFT':
            raise HTTPException(409, '盘点单已变化或已完成，请重新打开')
        check_date(payload.business_date)
        if payload.action == 'REFRESH':
            document['lines'] = count_snapshot(stock, [line['batch_id'] for line in document['lines']], entries)
            document['ledger_token'] = inventory_token(stock, entries)
            set_count_scope(document, stock, entries)
        elif payload.action in {'SAVE', 'POST'}:
            submitted = {line.batch_id: line for line in payload.lines}
            if len(submitted) != len(payload.lines) or set(submitted) != {line['batch_id'] for line in document['lines']}:
                raise HTTPException(422, '盘点明细与原单不一致，不能漏盘或添加未核对批次')
            for line in document['lines']:
                supplied = submitted[line['batch_id']]
                actual = ops.text(quantity(supplied.actual, '实盘数量', allow_zero=True)) if supplied.actual else ''
                if payload.action == 'POST' and actual == '':
                    raise HTTPException(422, '全部批次须填写实盘数量，零库存请明确填写0')
                line.update(actual=actual, reason=supplied.reason)
        if payload.action == 'POST':
            if not payload.confirmed:
                raise HTTPException(422, '请确认实盘数量、盘点截止及差异依据')
            changed = count_changes(document, stock, entries)
            if changed:
                raise HTTPException(409, f'盘点范围内有 {len(changed)} 个批次的记录发生变化，请刷新查看并更新账面重新盘点；本次未调整库存')
            if payload.business_date.isoformat() < document['business_date']:
                raise HTTPException(422, '盘点入账日期不能早于生成盘点单的日期')
            for index, line in enumerate(document['lines']):
                batch = ops.batch_for(stock, line['batch_id'])
                difference = Decimal(line['actual']) - Decimal(batch['quantity'])
                if difference and not line['reason']:
                    raise HTTPException(422, '盘盈或盘亏须逐行填写差异原因')
                # Zero differences also establish a dated audit dependency.
                body = dict(batch_id=batch['id'], quantity=ops.text(difference), actual_quantity=line['actual'], book_quantity=batch['quantity'],
                    stocktake_id=identifier, source_no=identifier, source_line=str(index+1), reason=line['reason'], allocations=[], counterparty=batch.get('counterparty',''))
                append(db, actor, stock, entries, 'STOCKTAKE', body, payload.business_date, uuid5(payload.request_id, f'count-{index}'), signature)
                line['difference'] = ops.text(difference)
            document['status'] = 'POSTED'
            document['posted_on'] = payload.business_date.isoformat()
        elif payload.action == 'CANCEL':
            if not payload.reason:
                raise HTTPException(422, '请填写取消原因')
            document['status'] = 'CANCELLED'
        document['revision'] += 1
        document['events'].append(dict(action=payload.action, actor_name=actor.display_name, at=business_now().isoformat(), reason=payload.reason))
        append(db, actor, stock, entries, 'STOCKTAKE_'+payload.action, {'result':document}, payload.business_date, payload.request_id, signature)
        state.revision += 1
        db.commit(); return document
    except Exception:
        db.rollback(); raise


def workspace(db):
    ops.ensure_schema(db)
    stock, entries = ops.projection(db, 'fabric')
    counts = list(count_documents(entries).values())
    for document in counts:
        document['changed_batch_ids'] = count_changes(document, stock, entries) if document['status'] == 'DRAFT' else []
        document['stale'] = bool(document['changed_batch_ids'])
        document['legacy_scope'] = document.get('count_scope') != 'BATCHES_V1'
    return dict(openings=[e['data']['result'] for e in reversed(list(entries.values())) if e['kind']=='OPENING_IMPORT'], counts=list(reversed(counts)))


def apply_inventory_entry(stock, entries, entry):
    """Replay only frozen facts; never recheck today's master or inventory token."""
    kind, body = entry['kind'], entry['data']
    if kind in META:
        entries[entry['id']] = entry
        return
    if kind == 'OPENING':
        ops.new_batch(stock, entry, body, quantity(body['quantity'], '期初数量'))
        stock[entry['id']].update(inbound_quantity='0', opening_quantity=body['quantity'], adjustment_quantity='0', inventory_origin='OPENING')
    else:
        batch = ops.batch_for(stock, body['batch_id'])
        if entry['business_date'] < batch.get('last_business_date', batch['received_on']):
            raise HTTPException(409, '盘点入账日期早于该批次最近操作')
        entry['batch_snapshot'] = copy.deepcopy(batch)
        ops.describe_item(body, batch)
        body['location'] = batch['location']
        delta = Decimal(body['quantity'])
        if Decimal(batch['quantity']) + delta < 0:
            raise HTTPException(409, '盘点差异将导致负库存')
        batch.update(quantity=ops.text(Decimal(batch['quantity'])+delta), adjustment_quantity=ops.text(Decimal(batch.get('adjustment_quantity','0'))+delta), last_business_date=entry['business_date'])
    entries[entry['id']] = entry
