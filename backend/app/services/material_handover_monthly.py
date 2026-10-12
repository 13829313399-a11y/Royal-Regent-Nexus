"""Live quantity reconciliation; reference prices never become approved payables."""
import calendar
import json
import re
from collections import defaultdict
from datetime import date
from decimal import Decimal
from fastapi import HTTPException
from sqlalchemy import select
from app.models.material_handover import HandoverDocument
from app.models.fabric_receiving import FabricReceipt
from app.services import material_handover as service
from app.services.fabric_procurement_parser import digest


def bounds(month):
    if not re.fullmatch(r'20[0-9]{2}-(0[1-9]|1[0-2])',month):raise HTTPException(422,'月份格式须为 YYYY-MM')
    year,number=map(int,month.split('-'))
    return date(year,number,1).isoformat(),date(year,number,calendar.monthrange(year,number)[1]).isoformat()


def report(db, access, month, supplier_id=''):
    service.ensure_schema(db)
    start,end=bounds(month)
    service.inventory.procurement.lock_factory(db)
    if access.issuer_id:
        if not access.supplier_id:raise HTTPException(403,'月度对账仅供已绑定基础资料的供应商查看')
        if supplier_id and supplier_id!=access.supplier_id:raise HTTPException(403,'只能查看本供应商对账')
        supplier_id=access.supplier_id
    rows=list(db.scalars(select(HandoverDocument).where(HandoverDocument.factory_id==service.FACTORY,HandoverDocument.status=='CONFIRMED')))
    docs={row.id:(row,json.loads(row.data_json)) for row in rows}
    entries=[service.inventory.entry_data(row) for row in service.inventory.documents(db,'fabric')]
    by_id={entry['id']:entry for entry in entries}
    events=defaultdict(list);unlinked=[]
    for entry in entries:
        if entry['kind']=='RECEIPT':source=entry;sign=1
        elif entry['kind']=='CORRECTION':
            source=by_id.get(entry['data'].get('original_id'))
            if not source or source['kind']!='RECEIPT':continue
            sign=-1
        else:continue  # workshop issues/returns/transfers are not supplier purchases
        body=source['data'];identifier=body.get('handover_document_id');line=body.get('handover_line_id')
        amount=Decimal(body['quantity'])*sign
        if identifier in docs and line:
            events[identifier,line].append((entry['business_date'],amount,entry['id']))
        elif start<=entry['business_date']<=end and not access.issuer_id:
            unlinked.append({'id':entry['id'],'source_no':body.get('source_no',''),'date':entry['business_date'],
                'counterparty':body.get('counterparty',''),'item_code':body.get('item_code',''),'quantity':str(amount),
                'unit':body.get('unit',''),'reason':'已有库存凭证，未关联供应商制单入口；不得重复入库'})
    if not access.issuer_id:
        for receipt in db.scalars(select(FabricReceipt).where(FabricReceipt.factory_id==service.FACTORY,
            FabricReceipt.receipt_date>=start,FabricReceipt.receipt_date<=end)):
            facts=json.loads(receipt.source_json)
            unlinked.append({'id':receipt.id,'source_no':receipt.delivery_reference,'date':receipt.receipt_date,
                'counterparty':facts.get('supplier',''),'item_code':facts.get('material_code',''),
                'quantity':receipt.quantity,'unit':receipt.unit,'reason':'已有采购实收原单，未关联本入口；不得重复入库'})
    details=[];groups={};units={}
    for identifier,(row,body) in docs.items():
        if body['kind']!='DELIVERY' or not body.get('document_date') or body['document_date']>end:continue
        party=service.supplier_identity(db,row,body)
        if supplier_id and party!=supplier_id:continue
        for line in body['lines']:
            target=Decimal(line['quantity']);moves=events[identifier,line['id']]
            prior=sum((amount for when,amount,_ in moves if when<start),Decimal(0))
            received=sum((amount for when,amount,_ in moves if start<=when<=end and amount>0),Decimal(0))
            reversed_amount=-sum((amount for when,amount,_ in moves if start<=when<=end and amount<0),Decimal(0))
            closing=prior+received-reversed_amount;outstanding=max(Decimal(0),target-closing)
            # Include outstanding carryovers and every movement (even a net-zero correction).
            if not (start<=body['document_date']<=end or received or reversed_amount or outstanding):continue
            notes=[]
            if not party:notes.append('供应商身份待关联')
            price=line.get('unit_price','')
            if not price:notes.append('单价未记录')
            elif not line.get('price_unit') or not line.get('currency'):notes.append('单价计价单位 / 币种待补')
            detail={'document_id':identifier,'line_id':line['id'],'source_no':body['source_no'],'document_date':body['document_date'],
                'supplier_id':party,'supplier_name':body['counterparty'],'item_code':line['item_code'],'item_name':line['item_name'],
                'specification':line.get('specification',''),'color':line.get('color',''),'unit':line['unit'],
                'declared':str(target),'prior_received':str(prior),'month_received':str(received),'month_reversed':str(reversed_amount),
                'month_net_received':str(received-reversed_amount),'cumulative_received':str(closing),'outstanding':str(outstanding),
                'unit_price':price,'price_unit':line.get('price_unit',''),'currency':line.get('currency',''),
                'price_note':line.get('price_note',''),'notes':notes}
            details.append(detail)
            key=(party or 'unbound:'+str(row.issuer_id or body['counterparty']),line['item_code'],line['unit'])
            group=groups.setdefault(key,{'supplier_id':party,'supplier_name':body['counterparty'],'item_code':line['item_code'],
                'item_name':line['item_name'],'unit':line['unit'],'month_received':Decimal(0),'month_reversed':Decimal(0),
                'month_net_received':Decimal(0),'outstanding':Decimal(0),'document_ids':set()})
            for field in ('month_received','month_reversed','month_net_received','outstanding'):group[field]+=Decimal(detail[field])
            group['document_ids'].add(identifier)
            unit=units.setdefault(line['unit'],{'unit':line['unit'],'month_received':Decimal(0),'month_reversed':Decimal(0),'month_net_received':Decimal(0)})
            for field in ('month_received','month_reversed','month_net_received'):unit[field]+=Decimal(detail[field])
    summaries=[{**group,**{k:str(group[k]) for k in ('month_received','month_reversed','month_net_received','outstanding')},
        'document_count':len(group['document_ids'])} for group in groups.values()]
    for group in summaries:group.pop('document_ids')
    totals=[{k:str(v) if isinstance(v,Decimal) else v for k,v in unit.items()} for unit in units.values()]
    summaries.sort(key=lambda row:(row['supplier_id'],row['item_code'],row['unit']))
    totals.sort(key=lambda row:row['unit'])
    unlinked.sort(key=lambda row:(row['date'],row['id']))
    details.sort(key=lambda row:(row['supplier_name'],row['document_date'],row['source_no'],row['line_id']))
    result={'month':month,'basis':'实际收货 / 冲销业务日期；不同单位分别汇总，领料不扣供应商进货账',
        'details':details,'summary':summaries,'totals':totals,'unlinked_receipts':unlinked,
        'closed':False,'amounts_calculated':False,'live':True}
    result['version']=digest(result)
    return result


def export_xlsx(report):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book=Workbook();sheet=book.active;sheet.title='单据数量对账'
    sheet.append([report['month'],'实时数量对账；价格仅供核对，不是关账或付款确认'])
    columns=[('supplier_name','供应商'),('source_no','送货单号'),('document_date','单据日期'),('item_code','物料编码'),
        ('item_name','名称'),('unit','库存单位'),('declared','相关原单申报量'),('prior_received','月初累计已收'),
        ('month_received','当月实收'),('month_reversed','当月冲销'),('month_net_received','当月净实收'),
        ('outstanding','月末未收'),('unit_price','参考单价'),('price_unit','计价单位'),('currency','币种'),('price_note','单价依据')]
    sheet.append([label for _,label in columns]+['待核对'])
    for row in report['details']:
        sheet.append([str(row.get(key,'')) for key,_ in columns]+['；'.join(row['notes'])])
    other=book.create_sheet('未关联入口的库存凭证')
    other.append(['原凭证','日期','往来单位','物料编码','数量','单位','说明'])
    for row in report['unlinked_receipts']:
        other.append([row[k] for k in ('source_no','date','counterparty','item_code','quantity','unit','reason')])
    for page in book:
        page.freeze_panes='A3' if page==sheet else 'A2'
        page.auto_filter.ref=page.dimensions
        for row in page:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'  # preserve codes and prevent formula injection
        for cell in page[2 if page==sheet else 1]:cell.font=Font(bold=True)
        from openpyxl.utils import get_column_letter
        for column in range(1,page.max_column+1):page.column_dimensions[get_column_letter(column)].width=20
    output=BytesIO();book.save(output);return output.getvalue()
