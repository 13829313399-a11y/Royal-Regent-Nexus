"""Read explicit Excel reference provenance without calculating workbook formulas."""
import io
import re


def links(filename, content, cells):
    if not filename.lower().endswith('.xlsx'):return {}
    import openpyxl
    from openpyxl.utils.cell import coordinate_to_tuple
    from openpyxl.utils import get_column_letter
    from app.services.carton_mark_ai_pilot import normalized
    native={c.location:c for c in cells};result={}
    book=openpyxl.load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
    try:
        # Only read already validated, selected cells. Never evaluate expressions,
        # follow external links or accept formula text as a business value.
        formulas={}
        for sheet in book:
            selected={loc.rsplit('!',1)[1]:c for loc,c in native.items() if loc.rsplit('!',1)[0]==sheet.title}
            if not selected:continue
            coordinates=[coordinate_to_tuple(co) for co in selected]
            height=max(r for r,c in coordinates);width=max(c for r,c in coordinates)
            if height>10000 or width>256 or height*width>200000:continue
            for row in sheet.iter_rows(max_row=height,max_col=width):
                for cell in row:
                    if cell.data_type=='f' and getattr(cell,'coordinate',None) in selected and isinstance(cell.value,str):formulas[f'{sheet.title}!{cell.coordinate}']=cell.value
        def reference(text,sheet):
            match=re.fullmatch(r"(?:(?:'((?:[^']|'')+)'|([^'!\[\]]+))!)?(\$?[A-Z]{1,3}\$?\d+)",text.strip(),re.I)
            if not match:return None
            return ((match[1].replace("''", "'") if match[1] else match[2]) or sheet)+'!'+match[3].replace('$','').upper()
        for location,formula in formulas.items():
            sheet=location.rsplit('!',1)[0];target=reference(formula.lstrip('='),sheet)
            if target is None:
                # A literal, exact-match VLOOKUP identifies a source column. Use
                # the existing displayed key and unique native row, not a formula
                # engine; arithmetic, approximate lookup and nested code are ignored.
                match=re.fullmatch(r'=VLOOKUP\(([^,]+),((?:\x27[^\x27]+\x27|[^!,]+)!)?(\$?[A-Z]{1,3})(?:\$?\d+)?:\$?([A-Z]{1,3})(?:\$?\d+)?,(\d+),(?:0|FALSE)\)',formula,re.I)
                if match:
                    key=native.get(reference(match[1],sheet));sh=(match[2] or (sheet+'!')).rstrip('!').strip("'")
                    first=openpyxl.utils.column_index_from_string(match[3].replace('$',''));index=int(match[5])
                    last=openpyxl.utils.column_index_from_string(match[4])
                    matches=[c for loc,c in native.items() if loc.rsplit('!',1)[0]==sh and coordinate_to_tuple(loc.rsplit('!',1)[1])[1]==first and key and normalized(c.text)==normalized(key.text)]
                    if len(matches)==1 and 1<=index<=last-first+1:
                        r,_=coordinate_to_tuple(matches[0].location.rsplit('!',1)[1]);target=f'{sh}!{get_column_letter(first+index-1)}{r}'
            source=native.get(location);destination=native.get(target)
            if source and destination and normalized(source.text)==normalized(destination.text):result[source.id]=destination.id
        # Resolve direct chains only, with a fixed depth and cycle protection.
        for key in list(result):
            seen={key};target=result[key]
            for _ in range(12):
                if target in seen:break
                seen.add(target)
                if target not in result:result[key]=target;break
                target=result[target]
    finally:book.close()
    return result
