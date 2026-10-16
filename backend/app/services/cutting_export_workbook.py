"""Bounded, lossless XLSX volumes for one authorized cutting report capture."""
from io import BytesIO
from hashlib import sha256
from html import escape
import json
import math
from zipfile import ZipFile, ZIP_STORED

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.page import PageMargins

from app.services import cutting_ops as c

MAX_ROWS = 20000
MAX_CELLS = 400000
MAX_TEXT_BYTES = 4 * 1024 * 1024
MAX_REPORT_CELLS = 2000000
MAX_VOLUMES = 100
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_PACKAGE_BYTES = 100 * 1024 * 1024


def checked(value):
    if isinstance(value, str):
        c.require(len(value) <= 32767, '单元格内容超长，不能静默截断导出', 413)
        c.require(not ILLEGAL_CHARACTERS_RE.search(value), '来源含不支持的控制字符，请核对原始依据', 422)
    return value if value is not None else '待核对'


def row_size(row):
    # Conservative XML payload allowance, including per-cell markup and escaped text.
    return sum(len(escape(str(checked(v))).encode('utf-8')) + 160 for v in row) + 100


def plan(report):
    """Stable row ranges: each logical data row occurs in exactly one volume."""
    tables = report['sheets']
    c.require(sum((len(t['rows']) + 1) * len(t['headers']) for t in tables) <= MAX_REPORT_CELLS,
              '完整报表超过两百万单元格的在线导出上限；未生成部分文件，请联系管理员安排完整离线导出', 413)
    overhead = sum(row_size(row) for row in report['metadata']) + 256000
    volumes, current = [], []
    cells, size = 0, overhead
    for index, table in enumerate(tables):
        width = len(table['headers'])
        header_size = row_size(table['headers'])
        start, count = 0, len(table['rows'])
        while start < count or (count == 0 and start == 0):
            end = start
            part_cells, part_size = width * 2, header_size * 2
            while end < count and end - start < MAX_ROWS:
                amount = row_size(table['rows'][end])
                if part_cells + width > MAX_CELLS - 10000 or part_size + amount > MAX_TEXT_BYTES - overhead:
                    break
                part_cells += width; part_size += amount; end += 1
            c.require(end > start or count == 0, '单行报表内容超过分卷容量；未生成部分文件，请联系管理员核对来源内容', 413)
            if current and (cells + part_cells > MAX_CELLS - 10000 or size + part_size > MAX_TEXT_BYTES):
                volumes.append(current); current = []; cells = 0; size = overhead
            current.append(dict(table=index, start=start, end=end))
            cells += part_cells; size += part_size
            start = end
            if count == 0: break
    if current: volumes.append(current)
    c.require(len(volumes) <= MAX_VOLUMES, '完整报表超过100卷的在线导出上限；未生成部分文件，请联系管理员安排完整离线导出', 413)
    return volumes or [[]]


def manifest(report, volumes):
    result, sequence = [], {}
    counts = {}
    for volume in volumes:
        for part in volume: counts[part['table']] = counts.get(part['table'], 0) + 1
    for number, volume in enumerate(volumes, 1):
        for part in volume:
            table = report['sheets'][part['table']]
            sequence[part['table']] = sequence.get(part['table'], 0) + 1
            name = table['name'] if counts[part['table']] == 1 else f"{table['name'][:24]}-{sequence[part['table']]:03d}"
            result.append(dict(volume=number, sheet=name, source=table['name'],
                               first=part['start'] + 1 if part['end'] else 0, last=part['end'],
                               rows=part['end'] - part['start'], total=len(table['rows'])))
    return result


def delivery(report):
    volumes = plan(report)
    return dict(format='xlsx' if len(volumes) == 1 else 'zip', volumes=len(volumes), manifest=manifest(report, volumes))


def print_summary(report):
    source = report['sheets'][0]
    rows = []
    if report['kind'] in {'period', 'payments'}:
        for values in source['rows']:
            rows.extend([[label, values[0], value] for label, value in zip(source['headers'][1:], values[1:])])
    else:
        for values in source['rows']:
            # Full vendor evidence stays in the electronic settlement summary; do not shrink it onto a cover.
            rows.extend([[label, values[4], value] for label, value in zip(source['headers'][:-1], values[:-1])])
    return dict(name='打印摘要', headers=['项目（引用汇总，不另累计）', '币种', '冻结值／当前余额值'], rows=rows)


def package(report, user, renderer):
    volumes = plan(report)
    catalog = manifest(report, volumes)
    manifest_table = dict(name='完整性清单', headers=['卷号', '工作表', '原表', '起始数据行', '结束数据行', '本段行数', '原表总行数'],
                          rows=[[r[k] for k in ('volume', 'sheet', 'source', 'first', 'last', 'rows', 'total')] for r in catalog])
    files = []
    for number, volume in enumerate(volumes, 1):
        entries = [r for r in catalog if r['volume'] == number]
        tables = [dict(name=r['sheet'], headers=report['sheets'][p['table']]['headers'],
                       rows=report['sheets'][p['table']]['rows'][p['start']:p['end']]) for p, r in zip(volume, entries)]
        # The summary is printed once only; metadata and coverage manifest repeat for traceability.
        if number == 1: tables.insert(0, print_summary(report))
        part_report = dict(report, sheets=[manifest_table] + tables,
            metadata=report['metadata'] + [['分卷', f'{number}/{len(volumes)}'],
                ['完整性说明', '明细数据行按清单连续分卷且不重复；打印摘要只在首卷，引用汇总不可再次累计。请保留全部卷。'],
                ['打印说明', '打印摘要适合整页打印；宽明细保留100%字号横向分页并重复首列、表头，长依据请在电子文件中查看。']])
        content = renderer(part_report, user)
        c.require(len(content) <= MAX_FILE_BYTES, '单卷文件超过20MB安全上限；整份下载已取消，请联系管理员核对报表容量', 413)
        files.append((f'cutting-{report["kind"]}-{report["month"]}-{report["token"][:12]}-{number:03d}.xlsx', content))
        c.require(sum(len(data) for _, data in files) <= MAX_PACKAGE_BYTES,
                  '完整报表超过100MB在线下载上限；未返回部分文件，请联系管理员安排完整离线导出', 413)
    if len(files) == 1: return files[0][1], 'xlsx'
    output = BytesIO()
    with ZipFile(output, 'w', ZIP_STORED) as archive:
        for name, content in files: archive.writestr(name, content)
        archive.writestr('manifest.json', json.dumps(dict(token=report['token'], month=report['month'],
            metadata=report['metadata'], coverage=catalog,
            files=[dict(name=name, sha256=sha256(content).hexdigest(), bytes=len(content)) for name, content in files]), ensure_ascii=False, indent=2))
    c.require(output.tell() <= MAX_PACKAGE_BYTES, '完整压缩包超过100MB在线上限；未返回部分文件，请联系管理员安排完整离线导出', 413)
    return output.getvalue(), 'zip'


def text_lines(value, width):
    # CJK glyphs occupy approximately two Excel character units. Preserve explicit newlines.
    return sum(max(1, math.ceil(sum(2 if ord(ch) > 255 else 1 for ch in line) / max(width - 2, 1)))
               for line in str(value).split('\n'))


def render_workbook(report, user, title):
    workbook = Workbook(); workbook.remove(workbook.active)
    metadata = report['metadata'] + [['导出人', user.display_name], ['导出人ID', user.id],
        ['取数／导出时间', report['captured_at']], ['报表依据校验值', report['token']]]
    warning = next((str(value) for key, value in metadata if key in {'使用状态', '口径'}), '')
    for table in [dict(name='报表说明', headers=['项目', '内容'], rows=metadata)] + report['sheets']:
        ws = workbook.create_sheet(table['name']); ws.sheet_view.showGridLines = False
        widths = [42, 98] if table['name'] == '报表说明' else [38, 12, 38] if table['name'] == '打印摘要' else [
            36 if any(k in h for k in ('依据', '凭据', '说明', '原因')) else 24 if 'ID' in h or '时间' in h else 18 for h in table['headers']]
        wide = sum(widths) > 155
        ws.append([title + ' · ' + warning]); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(widths))
        ws.append(table['headers'])
        for row in table['rows']: ws.append([checked(value) for value in row])
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str): cell.data_type = 's'
                cell.font = Font(name='Microsoft YaHei', size=10, bold=cell.row <= 2, color='FFFFFF' if cell.row <= 2 else '173E43')
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                if cell.row <= 2 or cell.row % 2 == 0: cell.fill = PatternFill('solid', fgColor='007F78' if cell.row <= 2 else 'EFF7F7')
            if row[0].row > 1:
                height = max(text_lines(cell.value or '', widths[i]) for i, cell in enumerate(row)) * 16 + 8
                ws.row_dimensions[row[0].row].height = min(300, max(32, height))
        for i, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = width
        ws.row_dimensions[1].height = 58
        ws.freeze_panes = 'B3'; ws.auto_filter.ref = f'A2:{get_column_letter(len(widths))}{max(ws.max_row, 2)}'
        ws.print_area = f'A1:{get_column_letter(len(widths))}{ws.max_row}'
        ws.print_title_rows = '1:2'; ws.print_options.horizontalCentered = True
        ws.page_margins = PageMargins(left=.3, right=.3, top=.5, bottom=.5, header=.2, footer=.2)
        ws.sheet_properties.pageSetUpPr.fitToPage = not wide
        ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A3
        if wide:
            ws.page_setup.scale = 100; ws.print_title_cols = 'A:A'
            width = widths[0]
            for index, value in enumerate(widths[1:], 1):
                if width + value > 150:
                    ws.col_breaks.append(Break(id=index)); width = widths[0]
                width += value
        else:
            ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.oddHeader.left.text = f'{report.get("month", "")} · {table["name"]}'
        ws.oddHeader.right.text = ('历史失效，仅供追溯' if '历史失效' in warning else '当前全期余额，非月末快照' if report['kind'] == 'payments' else '冻结版本')
        ws.oddFooter.center.text = '第 &P 页 / 共 &N 页'
        ws.oddFooter.right.text = f'华康 C 裁床部 · {report["token"][:12]}'
    stream = BytesIO(); workbook.save(stream)
    return stream.getvalue()
