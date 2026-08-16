from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.shared import Inches

from app.schemas.document_studio import DocumentSnapshot
from app.services.document_studio.contracts import DocumentExtractionRoute
from app.services.pdf_to_word import PdfToWordConversionError, PdfToWordResult
from app.services.pdf_to_word import convert_pdf_to_word as convert_editable_pdf_to_word


def _safe_output_name(source_filename: str) -> str:
    stem = re.sub(
        r'[\\/:*?"<>|]+',
        "_",
        Path(source_filename or "PDF文件.pdf").stem,
    ).strip(" .") or "PDF文件"
    return f"{stem}_版式保真.docx"


def convert_pdf_to_layout_preserving_word(
    pdf_bytes: bytes,
    source_filename: str,
) -> PdfToWordResult:
    try:
        pdf = pdfium.PdfDocument(pdf_bytes)
    except Exception as exc:
        raise PdfToWordConversionError("PDF 无法读取；请确认文件未加密、未损坏。") from exc
    page_count = len(pdf)
    if page_count < 1:
        raise PdfToWordConversionError("PDF 中没有可转换的页面。")
    if page_count > 80:
        raise PdfToWordConversionError("单次最多转换 80 页 PDF。")

    document = Document()
    document.core_properties.title = Path(source_filename or "PDF文件.pdf").name
    for page_index in range(page_count):
        page = pdf[page_index]
        width_points, height_points = page.get_size()
        section = document.sections[-1]
        section.page_width = Inches(width_points / 72)
        section.page_height = Inches(height_points / 72)
        section.orientation = (
            WD_ORIENT.LANDSCAPE
            if width_points > height_points
            else WD_ORIENT.PORTRAIT
        )
        section.top_margin = Inches(0)
        section.bottom_margin = Inches(0)
        section.left_margin = Inches(0)
        section.right_margin = Inches(0)
        rendered = page.render(scale=2).to_pil()
        image = BytesIO()
        rendered.save(image, format="PNG", optimize=True)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = 0
        paragraph.paragraph_format.space_after = 0
        run = paragraph.add_run()
        run.add_picture(image, width=Inches(width_points / 72))
        if page_index + 1 < page_count:
            run.add_break(WD_BREAK.PAGE)

    output = BytesIO()
    document.save(output)
    return PdfToWordResult(
        content=output.getvalue(),
        page_count=page_count,
        table_count=0,
        image_count=page_count,
        text_page_count=0,
        ocr_page_count=0,
        output_file_name=_safe_output_name(source_filename),
    )


def convert_pdf_to_word(
    pdf_bytes: bytes,
    source_filename: str,
    *,
    mode: str,
    snapshot: DocumentSnapshot | None = None,
) -> PdfToWordResult:
    if mode == "LAYOUT_PRESERVING":
        return convert_pdf_to_layout_preserving_word(pdf_bytes, source_filename)
    page_text_overrides = None
    if snapshot is not None:
        page_text_overrides = {
            page.page_number: tuple(
                (
                    float(block.bbox[1]),
                    float(block.bbox[3]),
                    block.normalized_text,
                )
                for block in page.blocks
                if block.normalized_text.strip()
            )
            for page in snapshot.pages
            if page.extraction_route == DocumentExtractionRoute.QWEN_OCR
            and page.blocks
        }
    return convert_editable_pdf_to_word(
        pdf_bytes,
        source_filename,
        page_text_overrides=page_text_overrides or None,
    )
