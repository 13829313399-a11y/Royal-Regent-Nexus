from __future__ import annotations

import re
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.schemas.document_studio import DocumentJobOptions, DocumentSnapshot
from app.services.pdf_to_excel import PdfToExcelResult, convert_pdf_to_excel


def _page_number(title: str) -> int | None:
    match = re.match(r"^第([1-9][0-9]*)页", title)
    return int(match.group(1)) if match else None


def _copy_values(source, target) -> None:
    for row in source.iter_rows(values_only=True):
        target.append(list(row))


def _apply_page_strategy(workbook) -> None:
    groups: dict[int, list] = {}
    for sheet in workbook.worksheets:
        number = _page_number(sheet.title)
        if number is not None:
            groups.setdefault(number, []).append(sheet)
    for page_number, sheets in groups.items():
        if len(sheets) < 2:
            continue
        target = workbook.create_sheet(f"第{page_number}页_汇总")
        for index, sheet in enumerate(sheets):
            if index:
                target.append([])
            target.append([sheet.title])
            _copy_values(sheet, target)
        for sheet in sheets:
            workbook.remove(sheet)


def _sheet_signature(sheet) -> tuple[str, ...] | None:
    for row in sheet.iter_rows(values_only=True):
        values = tuple("" if value is None else str(value).strip() for value in row)
        if any(values):
            return values
    return None


def _apply_merge_strategy(workbook) -> None:
    groups: dict[tuple[str, ...], list] = {}
    for sheet in list(workbook.worksheets):
        if sheet.title in {"转换说明", "来源映射"}:
            continue
        signature = _sheet_signature(sheet)
        if signature is not None and len(signature) > 1:
            groups.setdefault(signature, []).append(sheet)
    for sheets in groups.values():
        if len(sheets) < 2:
            continue
        target = sheets[0]
        for source in sheets[1:]:
            skipped_header = False
            for row in source.iter_rows(values_only=True):
                values = list(row)
                if not skipped_header and tuple(
                    "" if value is None else str(value).strip() for value in values
                ) == _sheet_signature(source):
                    skipped_header = True
                    continue
                target.append(values)
            workbook.remove(source)


def _smart_value(value: str):
    stripped = value.strip()
    if not stripped or (len(stripped) > 1 and stripped.startswith("0") and stripped.isdigit()):
        return value
    if re.fullmatch(r"[-+]?\d+", stripped):
        return int(stripped)
    if re.fullmatch(r"[-+]?(?:\d+\.\d+|\d+%)", stripped):
        try:
            if stripped.endswith("%"):
                return float(Decimal(stripped[:-1]) / 100)
            return float(Decimal(stripped))
        except InvalidOperation:
            return value
    try:
        return date.fromisoformat(stripped.replace("/", "-"))
    except ValueError:
        pass
    return value


def _apply_smart_types(workbook) -> None:
    for sheet in workbook.worksheets:
        if sheet.title in {"转换说明", "来源映射"}:
            continue
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    cell.value = _smart_value(cell.value)


def _append_source_mapping(workbook, snapshot: DocumentSnapshot) -> None:
    if "来源映射" in workbook.sheetnames:
        workbook.remove(workbook["来源映射"])
    sheet = workbook.create_sheet("来源映射")
    sheet.append(
        [
            "目标/证据 ID",
            "页码",
            "类型",
            "原始值",
            "规范值",
            "置信度",
            "来源",
            "坐标 (left, top, right, bottom)",
        ]
    )
    for page in snapshot.pages:
        for block in page.blocks:
            sheet.append(
                [
                    block.block_id,
                    page.page_number,
                    block.kind.value,
                    block.raw_text,
                    block.normalized_text,
                    block.confidence,
                    block.source.value,
                    ", ".join(f"{value:.2f}" for value in block.bbox),
                ]
            )
        for table in page.tables:
            for cell in table.cells:
                sheet.append(
                    [
                        cell.cell_id,
                        page.page_number,
                        cell.value_type.value,
                        cell.raw_text,
                        cell.normalized_value,
                        cell.confidence,
                        ", ".join(source.value for source in cell.sources),
                        ", ".join(f"{value:.2f}" for value in cell.source_bbox),
                    ]
                )
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = (28, 9, 20, 36, 36, 12, 24, 38)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def convert_pdf_to_evidence_workbook(
    pdf_bytes: bytes,
    source_filename: str,
    *,
    snapshot: DocumentSnapshot,
    options: DocumentJobOptions,
) -> PdfToExcelResult:
    converted = convert_pdf_to_excel(pdf_bytes, source_filename)
    workbook = load_workbook(BytesIO(converted.content))
    if options.sheet_strategy == "PAGE_PER_SHEET":
        _apply_page_strategy(workbook)
    elif options.sheet_strategy == "MERGE_SAME_SCHEMA":
        _apply_merge_strategy(workbook)
    if options.type_inference == "SMART":
        _apply_smart_types(workbook)
    _append_source_mapping(workbook, snapshot)
    summary = workbook["转换说明"]
    summary.append(["工作表策略", options.sheet_strategy])
    summary.append(["类型推断", options.type_inference])
    summary.append(["证据映射", "来源映射工作表记录页码、证据 ID、置信度与原始坐标。"])
    output = BytesIO()
    workbook.save(output)
    return PdfToExcelResult(
        content=output.getvalue(),
        page_count=converted.page_count,
        table_count=converted.table_count,
        text_page_count=converted.text_page_count,
        ocr_page_count=converted.ocr_page_count,
        output_file_name=converted.output_file_name,
    )
