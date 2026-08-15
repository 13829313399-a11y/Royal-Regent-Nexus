from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pypdfium2 as pdfium
from docx import Document
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFont

from app.schemas.document_studio import DocumentSnapshot


def _font(size: int = 22):
    candidates = (
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    )
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def _wrap(draw: ImageDraw.ImageDraw, text: str, font, width: int) -> str:
    lines: list[str] = []
    current = ""
    for char in text:
        candidate = current + char
        if current and draw.textbbox((0, 0), candidate, font=font)[2] > width:
            lines.append(current)
            current = char
        else:
            current = candidate
    if current:
        lines.append(current)
    return "\n".join(lines)


def _translated_page(page, *, scale: float = 2.0) -> Image.Image:
    image = Image.new(
        "RGB",
        (max(1, int(page.width * scale)), max(1, int(page.height * scale))),
        "white",
    )
    draw = ImageDraw.Draw(image)
    font = _font(max(12, int(10 * scale)))
    for block in page.blocks:
        left, top, right, _bottom = block.bbox
        draw.multiline_text(
            (int(left * scale), int(top * scale)),
            _wrap(
                draw,
                block.normalized_text,
                font,
                max(20, int((right - left) * scale)),
            ),
            fill="black",
            font=font,
            spacing=max(2, int(scale * 2)),
        )
    for table in page.tables:
        for cell in table.cells:
            left, top, right, _bottom = cell.source_bbox
            draw.multiline_text(
                (int(left * scale), int(top * scale)),
                _wrap(
                    draw,
                    cell.normalized_value,
                    font,
                    max(20, int((right - left) * scale)),
                ),
                fill="black",
                font=font,
                spacing=max(2, int(scale * 2)),
            )
    return image


def _source_pages(data: bytes, *, scale: float = 2.0) -> list[Image.Image]:
    document = pdfium.PdfDocument(data)
    try:
        return [
            document[index].render(scale=scale).to_pil().convert("RGB")
            for index in range(len(document))
        ]
    finally:
        document.close()


def _combine(source: Image.Image, translated: Image.Image, layout: str) -> Image.Image:
    if layout == "TRANSLATED_ONLY":
        return translated
    if layout == "SIDE_BY_SIDE":
        canvas = Image.new(
            "RGB",
            (source.width + translated.width, max(source.height, translated.height)),
            "white",
        )
        canvas.paste(source, (0, 0))
        canvas.paste(translated, (source.width, 0))
        return canvas
    canvas = Image.new(
        "RGB",
        (max(source.width, translated.width), source.height + translated.height),
        "white",
    )
    canvas.paste(source, (0, 0))
    canvas.paste(translated, (0, source.height))
    return canvas


def _pdf_bytes(images: list[Image.Image]) -> bytes:
    if not images:
        raise ValueError("translation renderer requires at least one page")
    output = BytesIO()
    images[0].save(
        output,
        format="PDF",
        save_all=True,
        append_images=images[1:],
        resolution=144,
    )
    return output.getvalue()


def _docx_bytes(snapshot: DocumentSnapshot) -> bytes:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.size = Pt(10)
    for page in snapshot.pages:
        document.add_heading(f"Page {page.page_number}", level=2)
        for block in page.blocks:
            document.add_paragraph(block.normalized_text)
        for source_table in page.tables:
            table = document.add_table(
                rows=source_table.row_count,
                cols=source_table.column_count,
            )
            table.style = "Table Grid"
            for cell in source_table.cells:
                match = re.search(r"-r([1-9][0-9]*)-c([1-9][0-9]*)$", cell.cell_id)
                if match is None:
                    continue
                row_index = int(match.group(1)) - 1
                column_index = int(match.group(2)) - 1
                if row_index < source_table.row_count and column_index < source_table.column_count:
                    table.cell(row_index, column_index).text = cell.normalized_value
        if page.page_number != snapshot.page_count:
            document.add_page_break()
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def render_pdf_translation(
    *,
    source_pdf: bytes,
    snapshot: DocumentSnapshot,
    source_filename: str,
    layout: str,
    include_editable_docx: bool,
) -> tuple[bytes, str, str]:
    sources = _source_pages(source_pdf)
    translated = [_translated_page(page) for page in snapshot.pages]
    rendered = [
        _combine(source, target, layout)
        for source, target in zip(sources, translated, strict=True)
    ]
    pdf = _pdf_bytes(rendered)
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(source_filename).stem).strip("._")
    stem = stem or "document"
    if not include_editable_docx:
        return pdf, f"{stem}_translated.pdf", "application/pdf"
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr(f"{stem}_translated.pdf", pdf)
        archive.writestr(f"{stem}_translated.docx", _docx_bytes(snapshot))
    return output.getvalue(), f"{stem}_translated.zip", "application/zip"
