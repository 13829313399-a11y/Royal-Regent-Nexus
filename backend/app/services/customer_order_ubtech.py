"""Huakang D UBTECH native purchase orders and explicit schedule extensions."""
from copy import copy
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, ROUND_CEILING
from io import BytesIO
import re

import pdfplumber


@dataclass(frozen=True)
class UbtechSpec:
    name: str = '优必选'
    po_extensions: tuple = ('.pdf',)
    input_template: str = 'HUAKANG_D_UBTECH_PURCHASE_ORDER_V1'


SPEC = UbtechSpec()
TARGET_TEMPLATE = 'HEYUAN_BUSINESS_UNIFIED_HUAKANG_D_UBTECH_V1'
EXTENSIONS = ('订单单价RMB', '订单单价HKD', '订单金额HKD', '订单单价USD', '订单金额USD', '客户跟单', '跟单')


def compact(value):
    return re.sub(r'\s+', '', value or '')


def parse_po(content, filename):
    from app.services.customer_order_unified import CustomerOrderUnifiedError as Error
    def number(value):
        try:
            result = Decimal(compact(value).replace(',', ''))
            if not result.is_finite() or result <= 0:
                raise ValueError()
            return result
        except Exception as exc:
            raise Error(f'{filename}：无效数量或金额 {value!r}') from exc

    with pdfplumber.open(BytesIO(content)) as pdf:
        text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
        if 'UBTECH' not in text.upper() or 'Purchase Order' not in text or '优必选' not in compact(text):
            raise Error(f'{filename}：不是可识别的优必选 UBTECH 文本型采购单')
        identities = set(re.findall(r'订单号\s*[:：]\s*(POORD\d+)', text, re.I))
        if len(identities) != 1:
            raise Error(f'{filename}：无法唯一识别优必选订单号')
        if not re.search(r'币别\s*[:：]\s*人民币', text):
            raise Error(f'{filename}：当前优必选规则仅适用于人民币 PO')
        po = next(iter(identities)).upper()
        dated = re.search(r'日期\s*[:：]\s*(\d{4}/\d{1,2}/\d{1,2})', text)
        order_date = dated[1] if dated else ''
        rows, totals, seen_lines = [], [], set()
        headers = None
        required = {'行号', '物料代码', '物料名称', '数量', '单位', '单价', '金额', '需求日期'}
        for page_no, page in enumerate(pdf.pages, 1):
            for table in page.extract_tables():
                for cells in table:
                    labels = [compact(c) for c in cells]
                    if required.issubset(labels):
                        headers = {label: i for i, label in enumerate(labels) if label}
                        continue
                    if headers is None:
                        continue
                    def get(field):
                        i = headers.get(field, len(cells))
                        return compact(cells[i]) if i < len(cells) else ''
                    if any(c == '合计' for c in labels):
                        totals.append((number(get('数量')), number(get('金额'))))
                        continue
                    if not any(labels):
                        continue
                    if not get('行号').isdigit() or not get('物料代码').isdigit():
                        raise Error(f'{filename} 第{page_no}页：存在无法识别的采购明细，不能跳过')
                    line = get('行号')
                    if line in seen_lines:
                        raise Error(f'{filename}：采购行号 {line} 重复，请核对PO页次')
                    seen_lines.add(line)
                    if get('单位').upper() not in {'PCS', 'PC', 'EA'}:
                        raise Error(f'{filename}：数量单位 {get("单位")} 不能按成品件数录入')
                    qty, price, amount = (number(get(k)) for k in ('数量', '单价', '金额'))
                    if qty != qty.to_integral_value() or abs(qty * price - amount) > Decimal('.01'):
                        raise Error(f'{filename} 第{line}行：数量或人民币金额核对不一致')
                    try:
                        y, m, d = (int(part) for part in get('需求日期').split('/'))
                        ship = date(y, m, d).isoformat()
                    except ValueError as exc:
                        raise Error(f'{filename} 第{line}行：需求日期无法识别') from exc
                    rows.append(dict(po_no=po, product_no=get('物料代码'), product_name_zh=get('物料名称'),
                                     quantity=qty, price_rmb=price, requested_ship_date=ship,
                                     po_date=order_date, page=page_no, line=line,
                                     standard='ROHS' if 'ROHS' in text.upper() else ''))
        if not rows or not totals:
            raise Error(f'{filename}：未找到完整采购明细及合计，请提供原始文本型PDF')
        total_qty = sum(r['quantity'] for r in rows)
        total_amount = sum(r['quantity'] * r['price_rmb'] for r in rows)
        if any(q != total_qty or abs(a - total_amount) > Decimal('.01') for q, a in totals):
            raise Error(f'{filename}：采购单总数量或总金额不一致，不能漏行导入')
    return rows


def create_rows(po_files, received_date, history):
    from app.services.customer_order_unified import _number_text, _unique_history_value, _issue
    rows = []
    for filename, content in po_files:
        for parsed in parse_po(content, filename):
            source = f'{filename} · 第{parsed["page"]}页第{parsed["line"]}行'
            name = _unique_history_value(history, parsed['product_no'], 'product_name_zh', SPEC.name)
            row = dict(id=f'ubtech-{len(rows)+1}', received_date=received_date,
                       po_no=parsed['po_no'], reference_no=parsed['po_no'], contract_no='', production_no='',
                       customer_name=SPEC.name, customer_country='优必选 / UBTECH', country='',
                       product_no=parsed['product_no'], product_name_zh=name or parsed['product_name_zh'], product_name_en='',
                       quantity=_number_text(parsed['quantity']), units_per_carton='', carton_count='',
                       standard=parsed['standard'], unit_price_hkd='', amount_hkd='', packaging='',
                       line_q='', customer_q='', requested_ship_date=parsed['requested_ship_date'],
                       manual='无', carton_mark='', date_code='',
                       input_template=SPEC.input_template, target_template=TARGET_TEMPLATE,
                       item_sheet_name='ITEM表', source_po_file_name=filename, _regional_customer='ubtech',
                       _ubtech_price_rmb=str(parsed['price_rmb']),
                       lineage={k: source for k in ('po_no', 'reference_no', 'product_no', 'quantity', 'requested_ship_date', 'standard')},
                       issues=[dict(_issue('warning', 'ubtech_ship_date_review', 'requested_ship_date',
                                      f'默认使用PO需求日期{parsed["requested_ship_date"]}，可按客户最新要求更正。', can_skip=True),
                                    can_edit=True, edit_field='requested_ship_date', edit_label='客要求走货期', edit_input_type='date'),
                               _issue('blocked', 'ubtech_factory_price_manual', 'unit_price_hkd',
                                      '出厂价需人工填写；订单换算价写入优必选专属列，不作为出厂价。', can_skip=True)])
            row['lineage'].update(received_date='本次填写的客户确认下单邮件日期',
                                  product_name_zh='当前排期同货号唯一历史名称' if name else source,
                                  unit_price_hkd='出厂价人工输入，不继承历史单价',
                                  order_price=f'PO人民币{parsed["price_rmb"]}；订单HKD=RMB/0.85，USD=HKD/7.75；保留计算精度',
                                  po_date=f'PO表头日期 {parsed["po_date"]}', manual='优必选教程：说明书无')
            rows.append(row)
    return rows, ['优必选订单换算价与手工出厂价分别记录；箱唛、生产单号和系统状态不从历史单复制。'], SPEC.input_template


def finish_row(row):
    from app.services.customer_order_unified import _number, _number_text
    row['target_template'] = TARGET_TEMPLATE
    qty, pack = _number(row.get('quantity')), _number(row.get('units_per_carton'))
    if qty is not None and pack and pack > 0:
        row['carton_count'] = _number_text((qty / pack).to_integral_value(rounding=ROUND_CEILING))


def write_summary(sheet, r, row, order_row):
    from app.services.customer_order_unified import _number, _excel_value
    if sheet.title == '接单表':
        sheet.cell(r, 14).value = _excel_value(_number(row.get('unit_price_hkd')))
    else:
        sheet.cell(r, 14).value = f'=IF(LEN(\'接单表\'!N{order_row})=0,"",\'接单表\'!N{order_row})'
    sheet.cell(r, 15).value = f'=IF(LEN(N{r})=0,"",ROUND(J{r}*N{r},2))'
    sheet.cell(r, 16).value = f'=IF(OR(LEN(N{r})=0,NOT(ISNUMBER($P$1)),$P$1<=0),"",N{r}/$P$1)'
    if sheet.title != '接单表':
        sheet.cell(r, 16).value = f'=IF(LEN(\'接单表\'!P{order_row})=0,"",\'接单表\'!P{order_row})'
    sheet.cell(r, 17).value = f'=IF(LEN(P{r})=0,"",ROUND(J{r}*P{r},2))'
    for c in range(14, 18):
        sheet.cell(r, c).number_format = '0.00'


def write_extensions(sheet, r, row):
    from openpyxl.utils import get_column_letter
    from app.services.customer_order_unified import CustomerOrderUnifiedError
    columns = {}
    for label in EXTENSIONS:
        found = [c.column for c in sheet[3] if c.column > 28 and compact(str(c.value or '')) == label]
        if len(found) > 1:
            raise CustomerOrderUnifiedError(f'优必选专属列 {label} 重复，无法确定写入位置')
        c = found[0] if found else sheet.max_column + 1
        if not found:
            sheet.cell(3, c)._style = copy(sheet.cell(3, 30)._style)
            sheet.cell(3, c).value = label
            sheet.column_dimensions[get_column_letter(c)].width = 18
        columns[label] = c
    ref = {k: f'{get_column_letter(c)}{r}' for k, c in columns.items()}
    values = [float(row['_ubtech_price_rmb']), f'={ref[EXTENSIONS[0]]}/0.85',
              f'=ROUND(K{r}*{ref[EXTENSIONS[1]]},2)', f'={ref[EXTENSIONS[1]]}/7.75',
              f'=ROUND(K{r}*{ref[EXTENSIONS[3]]},2)', 'UBTECH', '林乐谊']
    for (label, c), value in zip(columns.items(), values):
        cell = sheet.cell(r, c)
        cell._style = copy(sheet.cell(r, 30)._style)
        cell.value = value
        if c in [columns[k] for k in EXTENSIONS[:5]]:
            cell.number_format = '0.00'
