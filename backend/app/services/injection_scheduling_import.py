from __future__ import annotations

from collections import defaultdict
from dataclasses import asdict, dataclass
from datetime import date, datetime, time, timedelta, timezone
from hashlib import sha256
from io import BytesIO
import math
import re
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


PARSER_VERSION = "injection-scheduling-xlsx-v2"
MAX_BUSINESS_COLUMNS = 44
MAX_EMPTY_ROWS = 80
MACHINE_CLASSES = ("120A", "104A", "80A", "60A", "50A", "32A", "24A", "18A", "14A", "12A", "10A", "7A", "5A", "4A")
FORMULA_ERRORS = {"#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}


@dataclass(frozen=True)
class ImportIssue:
    severity: str
    code: str
    message: str
    sheet_name: str = ""
    source_row: int = 0
    field: str = ""
    source_value: str = ""


@dataclass(frozen=True)
class ParsedInjectionScheduleImport:
    parser_version: str
    factory_id: str
    source_file_name: str
    source_sha256: str
    source_size_bytes: int
    business_date: str
    normalized: dict[str, Any]
    summary: dict[str, Any]
    issues: tuple[ImportIssue, ...]

    @property
    def blocker_count(self) -> int:
        return sum(issue.severity == "blocker" for issue in self.issues)

    @property
    def warning_count(self) -> int:
        return sum(issue.severity == "warning" for issue in self.issues)

    def issues_as_dicts(self) -> list[dict[str, Any]]:
        return [asdict(issue) for issue in self.issues]


def _stable_id(prefix: str, factory_id: str, natural_key: str) -> str:
    digest = sha256(f"{factory_id}\0{natural_key}".encode("utf-8")).hexdigest()[:24]
    return f"{prefix}-{digest}"


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="seconds")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _identifier(cell: Any) -> str:
    value = getattr(cell, "value", cell)
    if value is None:
        return ""
    if isinstance(value, bool):
        return _text(value)
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        number_format = str(getattr(cell, "number_format", "") or "").split(";", 1)[0]
        if float(value).is_integer() and re.fullmatch(r"0+", number_format):
            return str(int(value)).zfill(len(number_format))
    return _text(value)


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if math.isfinite(float(value)):
            return float(value)
        return None
    match = re.search(r"-?\d+(?:\.\d+)?", _text(value).replace(",", ""))
    return float(match.group()) if match else None


def _integer(value: Any) -> int:
    parsed = _number(value)
    return max(0, int(round(parsed))) if parsed is not None else 0


def _optional_integer(value: Any) -> int | None:
    parsed = _number(value)
    return max(0, int(round(parsed))) if parsed is not None else None


def _iso_datetime(value: Any) -> str:
    if isinstance(value, datetime):
        aware = value.replace(tzinfo=timezone(timedelta(hours=8)))
        return aware.isoformat(timespec="seconds")
    if isinstance(value, date):
        aware = datetime.combine(value, time.min).replace(tzinfo=timezone(timedelta(hours=8)))
        return aware.isoformat(timespec="seconds")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            parsed = from_excel(value)
        except (TypeError, ValueError, OverflowError):
            parsed = None
        if isinstance(parsed, datetime):
            aware = parsed.replace(tzinfo=timezone(timedelta(hours=8)))
            return aware.isoformat(timespec="seconds")
        if isinstance(parsed, date):
            aware = datetime.combine(parsed, time.min).replace(tzinfo=timezone(timedelta(hours=8)))
            return aware.isoformat(timespec="seconds")
    text = _text(value)
    if not text:
        return ""
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            parsed = datetime.strptime(text, fmt).replace(tzinfo=timezone(timedelta(hours=8)))
            return parsed.isoformat(timespec="seconds")
        except ValueError:
            continue
    return ""


def _duration_days(value: Any) -> float | None:
    parsed = _number(value)
    if parsed is None:
        return None
    return round(parsed, 4)


def _excel_duration_hours(value: Any) -> float | None:
    parsed = _number(value)
    if parsed is None:
        return None
    return round(parsed * 24, 4)


def _machine_class(*values: Any) -> str:
    joined = " ".join(_text(value).upper() for value in values)
    for machine_class in MACHINE_CLASSES:
        if re.search(rf"(?<!\d){re.escape(machine_class)}(?!\d)", joined):
            return machine_class
    return ""


def _tonnage(*values: Any) -> int | None:
    joined = " ".join(_text(value).upper() for value in values)
    match = re.search(r"(\d+(?:\.\d+)?)\s*T(?:\b|$)", joined)
    if match:
        return int(round(float(match.group(1))))
    for value in values:
        parsed = _number(value)
        if parsed is not None and 40 <= parsed <= 3000:
            return int(round(parsed))
    return None


def _shot_capacity(value: Any) -> float | None:
    match = re.search(r"(\d+(?:\.\d+)?)\s*G(?:\b|/)", _text(value).upper())
    return float(match.group(1)) if match else None


def _dimensions(value: Any) -> tuple[float | None, float | None]:
    match = re.search(
        r"(\d+(?:\.\d+)?)\s*(?:MM)?\s*[×X*]\s*(\d+(?:\.\d+)?)",
        _text(value).upper(),
    )
    if not match:
        return None, None
    return float(match.group(1)), float(match.group(2))


def _robot_level(value: Any) -> str:
    normalized = _text(value)
    if any(token in normalized for token in ("双臂", "五轴", "多轴")):
        return "multi-arm"
    if any(token in normalized for token in ("单臂", "三轴")):
        return "single-arm"
    return "none"


def _machine_kind(value: Any) -> str:
    normalized = _text(value)
    if "全电" in normalized:
        return "全电动"
    if "高速" in normalized:
        return "高速机"
    return "普通机"


def _machine_state(*values: Any) -> str:
    joined = " ".join(_text(value) for value in values)
    if any(token in joined for token in ("机故", "故障", "维修")):
        return "fault"
    if any(token in joined for token in ("保养", "停机")):
        return "maintenance"
    if any(token in joined for token in ("特急", "紧急")):
        return "urgent"
    return "idle"


def _priority(*values: Any) -> tuple[str, str]:
    joined = " ".join(_text(value) for value in values if _text(value))
    if "特急" in joined:
        return "P0", joined
    if "急" in joined:
        return "P1", joined
    return "P2", joined


def _color_family(value: Any) -> str:
    normalized = _text(value).lower()
    if any(token in normalized for token in ("透明", "自然", "本色")):
        return "natural"
    if "黑" in normalized:
        return "black"
    if any(token in normalized for token in ("白", "浅", "粉", "黄")):
        return "light"
    if any(token in normalized for token in ("深", "蓝", "紫", "绿", "红")):
        return "dark"
    if normalized:
        return "medium"
    return "special"


def _duration_hours(remaining: int, daily_target: int) -> float:
    if remaining <= 0:
        return 0.0
    if daily_target <= 0:
        return 24.0
    return round(max(0.25, remaining / daily_target * 24), 2)


def _iter_business_rows(sheet, *, start_row: int = 1):
    empty_streak = 0
    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=start_row,
            max_col=MAX_BUSINESS_COLUMNS,
            values_only=True,
        ),
        start=start_row,
    ):
        if any(value not in (None, "") for value in row):
            empty_streak = 0
            yield row_number, row
        else:
            empty_streak += 1
            if empty_streak >= MAX_EMPTY_ROWS:
                break


def _iter_business_cell_rows(cached_sheet, formula_sheet, *, start_row: int = 1):
    empty_streak = 0
    cached_rows = cached_sheet.iter_rows(
        min_row=start_row,
        max_col=MAX_BUSINESS_COLUMNS,
        values_only=False,
    )
    formula_rows = formula_sheet.iter_rows(
        min_row=start_row,
        max_col=MAX_BUSINESS_COLUMNS,
        values_only=False,
    )
    for row_number, (cached_cells, formula_cells) in enumerate(
        zip(cached_rows, formula_rows, strict=False),
        start=start_row,
    ):
        cached_values = tuple(cell.value for cell in cached_cells)
        formula_values = tuple(cell.value for cell in formula_cells)
        if any(value not in (None, "") for value in (*cached_values, *formula_values)):
            empty_streak = 0
            yield row_number, cached_cells, cached_values, formula_values
        else:
            empty_streak += 1
            if empty_streak >= MAX_EMPTY_ROWS:
                break


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=") or value.__class__.__name__ == "ArrayFormula"


def _append_formula_issues(
    *,
    row_number: int,
    sheet_name: str,
    cached_values: tuple[Any, ...],
    formula_values: tuple[Any, ...],
    columns: dict[str, int],
    issues: list[ImportIssue],
) -> None:
    checked_columns: set[int] = set()
    for field, column in columns.items():
        if column in checked_columns:
            continue
        checked_columns.add(column)
        formula_value = formula_values[column]
        if not _is_formula(formula_value):
            continue
        cached_value = cached_values[column]
        if isinstance(cached_value, str) and cached_value.strip().upper() in FORMULA_ERRORS:
            issues.append(
                ImportIssue(
                    severity="warning",
                    code="formula_error",
                    message=f"{field} 的公式结果为 {cached_value}，未按 0 导入",
                    sheet_name=sheet_name,
                    source_row=row_number,
                    field=field,
                    source_value=_text(cached_value),
                )
            )
        elif cached_value is None:
            issues.append(
                ImportIssue(
                    severity="warning",
                    code="formula_cache_missing",
                    message=f"{field} 缺少已计算公式缓存，未按 0 导入",
                    sheet_name=sheet_name,
                    source_row=row_number,
                    field=field,
                    source_value=_text(formula_value),
                )
            )


def _detect_template(sheet_names: set[str], factory_id: str) -> str:
    if "排期表" in sheet_names and "厂区现有啤机" in sheet_names:
        if factory_id != "huakang-b":
            raise ValueError("该工作簿属于华康B模板，只能导入华康B厂区")
        return "huakang-b"
    if "计划表" in sheet_names and "机台完成时间" in sheet_names:
        if factory_id != "huaxing":
            raise ValueError("该工作簿属于华兴模板，只能导入华兴厂区")
        return "huaxing"
    raise ValueError("无法识别注塑排产工作簿模板")


def _parse_huakang_machines(workbook, factory_id: str, issues: list[ImportIssue]) -> list[dict[str, Any]]:
    sheet = workbook["厂区现有啤机"]
    machines: list[dict[str, Any]] = []
    for row_number, row in _iter_business_rows(sheet, start_row=4):
        code = _text(row[0])
        name = _text(row[1])
        if not code or not name:
            continue
        machine_class = _machine_class(row[3])
        capacity = _shot_capacity(row[5])
        width, height = _dimensions(row[6])
        tonnage = _tonnage(row[7])
        robot = _robot_level(row[9])
        missing = [
            label
            for label, present in (
                ("机安", bool(machine_class)),
                ("射胶量", capacity is not None),
                ("容模宽高", width is not None and height is not None),
                ("吨位", tonnage is not None),
            )
            if not present
        ]
        if missing:
            issues.append(
                ImportIssue(
                    severity="warning",
                    code="machine_capability_incomplete",
                    message=f"机台 {code} 缺少能力参数：{'、'.join(missing)}",
                    sheet_name=sheet.title,
                    source_row=row_number,
                    field="machine_capability",
                    source_value=" | ".join(_text(value) for value in row[:10]),
                )
            )
        machine_id = _stable_id("ism", factory_id, code)
        machines.append(
            {
                "id": machine_id,
                "machineCode": code,
                "machineName": f"{code}号机",
                "workshop": "华康B啤机部",
                "state": _machine_state(row[18]),
                "kind": _machine_kind(row[8]),
                "machineClass": machine_class,
                "tonnage": tonnage,
                "injectionCapacityG": capacity,
                "safetyUtilization": 0.82,
                "moldWidthMm": width,
                "moldHeightMm": height,
                "robotLevel": robot,
                "capabilities": {
                    "supportsCorePull": False,
                    "supportsUnscrewing": False,
                    "compatibleMaterials": [],
                    "screwType": "standard",
                    "machineModel": _text(row[2]),
                    "robotModel": _text(row[10]),
                },
                "restrictions": [],
                "completeness": "complete" if not missing else "needs_review",
                "sourceLineage": {"sheet": sheet.title, "row": row_number},
            }
        )
    return machines


def _parse_huaxing_machines(workbook, factory_id: str, issues: list[ImportIssue]) -> list[dict[str, Any]]:
    sheet = workbook["机台完成时间"]
    machines: list[dict[str, Any]] = []
    for row_number, row in _iter_business_rows(sheet, start_row=3):
        code = _text(row[0])
        descriptor = _text(row[2])
        if not code or not descriptor:
            continue
        machine_class = _machine_class(descriptor)
        tonnage = _tonnage(descriptor)
        note = " ".join(_text(value) for value in row[5:8] if _text(value))
        missing = ["射胶量", "容模宽高"]
        if not machine_class:
            missing.insert(0, "机安")
        issues.append(
            ImportIssue(
                severity="warning",
                code="machine_capability_incomplete",
                message=f"机台 {code} 缺少能力参数：{'、'.join(missing)}，需人工复核",
                sheet_name=sheet.title,
                source_row=row_number,
                field="machine_capability",
                source_value=descriptor,
            )
        )
        machines.append(
            {
                "id": _stable_id("ism", factory_id, code),
                "machineCode": code,
                "machineName": f"{code}号机",
                "workshop": "华兴啤机部",
                "state": _machine_state(note),
                "kind": _machine_kind(row[3]),
                "machineClass": machine_class,
                "tonnage": tonnage,
                "injectionCapacityG": None,
                "safetyUtilization": 0.82,
                "moldWidthMm": None,
                "moldHeightMm": None,
                "robotLevel": _robot_level(row[4]),
                "capabilities": {
                    "supportsCorePull": False,
                    "supportsUnscrewing": False,
                    "compatibleMaterials": [],
                    "screwType": "PVC" if "PVC" in note.upper() else "standard",
                    "machineModel": descriptor,
                    "robotModel": _text(row[4]),
                },
                "restrictions": [note] if note else [],
                "completeness": "needs_review",
                "sourceLineage": {"sheet": sheet.title, "row": row_number},
            }
        )
    return machines


def _parse_huaxing_mold_master(workbook, factory_id: str) -> dict[str, dict[str, Any]]:
    if "机安" not in workbook.sheetnames:
        return {}
    sheet = workbook["机安"]
    result: dict[str, dict[str, Any]] = {}
    for row_number, row in _iter_business_rows(sheet, start_row=2):
        mold_no = _text(row[0])
        if not mold_no:
            continue
        result[mold_no.casefold()] = {
            "id": _stable_id("ismold", factory_id, mold_no),
            "moldNo": mold_no,
            "productName": _text(row[1]),
            "status": "available",
            "machineClassRequirement": _machine_class(row[6]),
            "lengthMm": _number(row[19]),
            "widthMm": None,
            "thicknessMm": None,
            "shotWeightG": _number(row[12]),
            "material": _text(row[11]),
            "robotRequirement": _robot_level(row[7]),
            "corePullRequired": "潜水" in _text(row[13]) and "否" not in _text(row[13]),
            "unscrewRequired": "牙" in " ".join(_text(value) for value in row[13:17]),
            "attributes": {
                "color": _text(row[9]),
                "automation": _text(row[15]),
                "suggestion": _text(row[16]),
                "remark": _text(row[14]),
            },
            "completeness": "needs_review",
            "sourceLineage": {"sheet": sheet.title, "row": row_number},
        }
    return result


def _task_columns(template: str) -> tuple[str, dict[str, int]]:
    if template == "huakang-b":
        return "排期表", {
            "machine": 1,
            "automation": 3,
            "remark": 3,
            "item": 4,
            "machine_class": 5,
            "mold": 6,
            "product": 7,
            "order": 8,
            "warehouse": 9,
            "set_quantity": 10,
            "order_qty": 11,
            "completed": 12,
            "remaining": 13,
            "target": 14,
            "material": 15,
            "water_ratio": 16,
            "color": 17,
            "color_powder": 18,
            "net_weight": 19,
            "gross_weight": 20,
            "material_weight": 21,
            "unit_price": 22,
            "order_date": 23,
            "delivery_start": 24,
            "delivery_due": 25,
            "mold_change_reference": 26,
            "color_change_reference": 27,
            "changeover": 28,
            "downtime": 29,
            "planned_start": 30,
            "planned_end": 31,
            "planned_month": 32,
            "inbound": 33,
            "slack": 34,
            "production_days": 35,
            "paint": 36,
            "shift_end": 37,
            "shift_time": 38,
            "shift_target": 39,
            "material_shortage": 40,
            "allocated_material": 41,
            "day_shift": 42,
            "night_shift": 43,
        }
    return "计划表", {
        "machine": 1,
        "automation": 3,
        "remark": 4,
        "machine_class": 5,
        "mold": 6,
        "product": 7,
        "order": 8,
        "item": 9,
        "set_quantity": 10,
        "order_qty": 11,
        "completed": 12,
        "remaining": 13,
        "target": 14,
        "water_ratio": 15,
        "color": 16,
        "color_powder": 17,
        "material": 18,
        "net_weight": 19,
        "gross_weight": 20,
        "material_weight": 21,
        "unit_price": 22,
        "outsourcing_unit_price": 23,
        "ratio": 24,
        "order_date": 25,
        "delivery_start": 26,
        "delivery_due": 27,
        "mold_change_reference": 28,
        "color_change_reference": 29,
        "changeover": 30,
        "downtime": 31,
        "planned_start": 32,
        "planned_end": 33,
        "planned_month": 34,
        "inbound": 35,
        "slack": 36,
        "paint": 37,
        "production_days": 38,
        "shift_end": 39,
        "shift_time": 40,
        "shift_target": 41,
        "secondary_machine_class": 42,
        "warehouse": 43,
    }


def _parse_plan_rows(
    workbook,
    formula_workbook,
    factory_id: str,
    template: str,
    machines: list[dict[str, Any]],
    issues: list[ImportIssue],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    sheet_name, columns = _task_columns(template)
    sheet = workbook[sheet_name]
    formula_sheet = formula_workbook[sheet_name]
    machine_by_code = {machine["machineCode"].casefold(): machine for machine in machines}
    mold_master = _parse_huaxing_mold_master(workbook, factory_id) if template == "huaxing" else {}
    molds: dict[str, dict[str, Any]] = dict(mold_master)
    orders: dict[str, dict[str, Any]] = {}
    tasks: list[dict[str, Any]] = []
    sequence_by_machine: defaultdict[str, int] = defaultdict(int)
    for row_number, cells, row, formula_row in _iter_business_cell_rows(
        sheet,
        formula_sheet,
        start_row=4,
    ):
        machine_cell = _identifier(cells[columns["machine"]])
        mold_no = _identifier(cells[columns["mold"]])
        order_no = _identifier(cells[columns["order"]])
        remaining = _integer(row[columns["remaining"]])
        is_machine_header = (
            _text(row[0])
            and _text(row[0]).casefold() == machine_cell.casefold()
            and bool(mold_no)
        )
        if is_machine_header:
            continue
        if mold_no and order_no:
            _append_formula_issues(
                row_number=row_number,
                sheet_name=sheet.title,
                cached_values=row,
                formula_values=formula_row,
                columns=columns,
                issues=issues,
            )
        if remaining <= 0 or not mold_no or not order_no:
            continue
        machine = machine_by_code.get(machine_cell.casefold())

        product_name = _text(row[columns["product"]])
        item_no = _identifier(cells[columns["item"]])
        warehouse = _identifier(cells[columns["warehouse"]])
        material = _text(row[columns["material"]])
        color = _text(row[columns["color"]])
        machine_class_requirement = _machine_class(row[columns["machine_class"]])
        mold_key = mold_no.casefold()
        mold = molds.get(mold_key)
        if mold is None:
            shot_weight = _number(row[columns["gross_weight"]]) or _number(row[columns["net_weight"]])
            mold = {
                "id": _stable_id("ismold", factory_id, mold_no),
                "moldNo": mold_no,
                "productName": product_name,
                "status": "available",
                "machineClassRequirement": machine_class_requirement,
                "lengthMm": None,
                "widthMm": None,
                "thicknessMm": None,
                "shotWeightG": shot_weight,
                "material": material,
                "robotRequirement": _robot_level(row[columns.get("automation", 0)]),
                "corePullRequired": False,
                "unscrewRequired": "牙" in _text(row[columns["remark"]]),
                "attributes": {
                    "color": color,
                    "paint": _text(row[columns["paint"]]),
                },
                "completeness": "needs_review",
                "sourceLineage": {"sheet": sheet.title, "row": row_number},
            }
            molds[mold_key] = mold
        else:
            if not mold.get("productName"):
                mold["productName"] = product_name
            if not mold.get("machineClassRequirement"):
                mold["machineClassRequirement"] = machine_class_requirement
            if not mold.get("material"):
                mold["material"] = material

        if not mold.get("shotWeightG"):
            issues.append(
                ImportIssue(
                    severity="warning",
                    code="mold_shot_weight_missing",
                    message=f"模具 {mold_no} 缺少单啤重量，候选机校验需人工复核",
                    sheet_name=sheet.title,
                    source_row=row_number,
                    field="shot_weight",
                    source_value=_text(row[columns["gross_weight"]]),
                )
            )

        natural_key = "|".join((order_no, item_no, mold_no))
        priority_code, priority_flag = _priority(
            row[columns["remark"]],
            row[columns["paint"]],
        )
        slack_days = _duration_days(row[columns["slack"]])
        worksheet = {
            "automationMode": (
                _text(row[columns["automation"]])
                if template == "huaxing"
                else ""
            ),
            "remark": _text(row[columns["remark"]]),
            "warehouse": warehouse,
            "machineClassRequirement": machine_class_requirement,
            "setQuantity": _optional_integer(row[columns["set_quantity"]]),
            "waterRatio": _text(row[columns["water_ratio"]]),
            "colorPowder": _identifier(cells[columns["color_powder"]]),
            "netWeightGrams": _number(row[columns["net_weight"]]),
            "grossWeightGrams": _number(row[columns["gross_weight"]]),
            "materialWeightKg": _number(row[columns["material_weight"]]),
            "orderDate": _iso_datetime(row[columns["order_date"]]),
            "deliveryStartAt": _iso_datetime(row[columns["delivery_start"]]),
            "deliveryDueAt": _iso_datetime(row[columns["delivery_due"]]),
            "moldChangeReferenceHours": _excel_duration_hours(
                row[columns["mold_change_reference"]]
            ),
            "colorChangeReferenceHours": _excel_duration_hours(
                row[columns["color_change_reference"]]
            ),
            "changeoverHours": _excel_duration_hours(row[columns["changeover"]]),
            "downtimeHours": _excel_duration_hours(row[columns["downtime"]]),
            "plannedProductionAt": _iso_datetime(row[columns["planned_start"]]),
            "plannedCompletionAt": _iso_datetime(row[columns["planned_end"]]),
            "plannedCompletionMonth": _text(row[columns["planned_month"]]),
            "inboundAt": _iso_datetime(row[columns["inbound"]]),
            "deliverySlackDays": slack_days,
            "sprayPaint": _text(row[columns["paint"]]),
            "productionDays": _duration_days(row[columns["production_days"]]),
            "materialShortage": (
                _number(row[columns["material_shortage"]])
                if "material_shortage" in columns
                else None
            ),
            "allocatedMaterialQuantity": (
                _number(row[columns["allocated_material"]])
                if "allocated_material" in columns
                else None
            ),
            "shiftEndAt": _text(row[columns["shift_end"]]),
            "shiftTime": _text(row[columns["shift_time"]]),
            "shiftTarget": _optional_integer(row[columns["shift_target"]]),
            "dayShiftQuantity": (
                _optional_integer(row[columns["day_shift"]])
                if "day_shift" in columns
                else None
            ),
            "nightShiftQuantity": (
                _optional_integer(row[columns["night_shift"]])
                if "night_shift" in columns
                else None
            ),
            "sourceSheet": sheet.title,
            "sourceRow": row_number,
        }
        restricted_worksheet = {
            "unitPricePerShot": _number(row[columns["unit_price"]]),
            "outsourcingUnitPrice": (
                _number(row[columns["outsourcing_unit_price"]])
                if "outsourcing_unit_price" in columns
                else None
            ),
            "ratio": (
                _number(row[columns["ratio"]])
                if "ratio" in columns
                else None
            ),
            "sourceSheet": sheet.title,
            "sourceRow": row_number,
        }
        order = {
            "id": _stable_id("isorder", factory_id, natural_key),
            "naturalKey": natural_key,
            "orderNo": order_no,
            "itemNo": item_no,
            "moldId": mold["id"],
            "productName": product_name,
            "orderQty": _integer(row[columns["order_qty"]]),
            "completedQty": _integer(row[columns["completed"]]),
            "remainingQty": remaining,
            "dailyTarget": _integer(row[columns["target"]]),
            "material": material,
            "color": color,
            "deliveryDueAt": _iso_datetime(row[columns["delivery_due"]]),
            "priorityCode": priority_code,
            "priorityFlag": priority_flag,
            "requirement": {
                "colorFamily": _color_family(color),
                "worksheet": worksheet,
                "restrictedWorksheet": restricted_worksheet,
            },
            "completeness": (
                "complete"
                if machine_class_requirement and material and mold.get("shotWeightG")
                else "needs_review"
            ),
            "sourceLineage": {"sheet": sheet.title, "row": row_number},
        }
        orders[natural_key] = order

        # Rows without a recognized machine belong to the workbook's waiting
        # pool. They remain in the imported order master but are not silently
        # assigned to the preceding machine.
        if machine is None:
            continue

        sequence = sequence_by_machine[machine["id"]]
        sequence_by_machine[machine["id"]] += 1
        duration = _duration_hours(order["remainingQty"], order["dailyTarget"])
        planned_start = _iso_datetime(row[columns["planned_start"]])
        planned_end = _iso_datetime(row[columns["planned_end"]])
        slack_value = slack_days * 24 if slack_days is not None else None
        risk = "urgent" if priority_code == "P0" else "normal"
        if slack_value is not None and slack_value < 0:
            risk = "overdue"
        elif order["completeness"] != "complete":
            risk = "incomplete"
        tasks.append(
            {
                "id": _stable_id("istask", factory_id, f"{natural_key}|{row_number}"),
                "machineId": machine["id"],
                "orderId": order["id"],
                "sequence": sequence,
                "current": sequence == 0,
                "locked": sequence == 0,
                "plannedStart": planned_start,
                "plannedEnd": planned_end,
                "risk": risk,
                "productionDurationHours": duration,
                "inboundAt": _iso_datetime(row[columns["inbound"]]),
                "slackHours": slack_value if slack_value is not None else 0,
                "remark": _text(row[columns["remark"]]),
                "worksheet": worksheet,
                "sourceLineage": {"sheet": sheet.title, "row": row_number},
            }
        )

    return list(molds.values()), list(orders.values()), tasks


def parse_injection_schedule_workbook(
    content: bytes,
    *,
    factory_id: str,
    source_file_name: str,
    business_date: str,
) -> ParsedInjectionScheduleImport:
    if not content:
        raise ValueError("上传文件为空")
    try:
        parsed_business_date = date.fromisoformat(business_date)
    except ValueError as exc:
        raise ValueError("business_date 必须为 YYYY-MM-DD") from exc

    workbook = None
    formula_workbook = None
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        formula_workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        if workbook is not None:
            workbook.close()
        if formula_workbook is not None:
            formula_workbook.close()
        raise ValueError("无法读取 Excel 工作簿，请确认文件未损坏且格式为 xlsx") from exc

    issues: list[ImportIssue] = []
    try:
        template = _detect_template(set(workbook.sheetnames), factory_id)
        machines = (
            _parse_huakang_machines(workbook, factory_id, issues)
            if template == "huakang-b"
            else _parse_huaxing_machines(workbook, factory_id, issues)
        )
        molds, orders, tasks = _parse_plan_rows(
            workbook,
            formula_workbook,
            factory_id,
            template,
            machines,
            issues,
        )
    finally:
        if workbook is not None:
            workbook.close()
        if formula_workbook is not None:
            formula_workbook.close()

    if not machines:
        issues.append(
            ImportIssue(
                severity="blocker",
                code="machines_missing",
                message="工作簿中没有可识别的机台资料",
            )
        )
    if not tasks:
        issues.append(
            ImportIssue(
                severity="blocker",
                code="active_tasks_missing",
                message="工作簿中没有剩余数量大于 0 的有效排产任务",
            )
        )

    anchor_at = datetime.combine(parsed_business_date, time(hour=8)).replace(
        tzinfo=timezone(timedelta(hours=8))
    )
    normalized = {
        "factoryId": factory_id,
        "template": template,
        "businessDate": parsed_business_date.isoformat(),
        "anchorAt": anchor_at.isoformat(timespec="seconds"),
        "sourceLabel": source_file_name,
        "machines": machines,
        "molds": molds,
        "orders": orders,
        "tasks": tasks,
    }
    summary = {
        "template": template,
        "machineCount": len(machines),
        "moldCount": len(molds),
        "orderCount": len(orders),
        "taskCount": len(tasks),
        "currentTaskCount": sum(bool(task["current"]) for task in tasks),
        "blockerCount": sum(issue.severity == "blocker" for issue in issues),
        "warningCount": sum(issue.severity == "warning" for issue in issues),
        "canConfirm": not any(issue.severity == "blocker" for issue in issues),
    }
    return ParsedInjectionScheduleImport(
        parser_version=PARSER_VERSION,
        factory_id=factory_id,
        source_file_name=source_file_name,
        source_sha256=sha256(content).hexdigest(),
        source_size_bytes=len(content),
        business_date=parsed_business_date.isoformat(),
        normalized=normalized,
        summary=summary,
        issues=tuple(issues),
    )
