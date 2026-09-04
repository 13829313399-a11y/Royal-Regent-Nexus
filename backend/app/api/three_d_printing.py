from __future__ import annotations

from datetime import UTC, datetime
from hmac import compare_digest
from io import BytesIO
from typing import Annotated

from fastapi import (
    APIRouter,
    Depends,
    File,
    Header,
    HTTPException,
    Query,
    Request,
    UploadFile,
)
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.db import get_db
from app.models.three_d_printing import (
    ThreeDPrintingPrinter,
    ThreeDPrintingProduct,
    ThreeDPrintingProductImage,
)
from app.schemas.three_d_printing import (
    ThreeDAuditEventOut,
    ThreeDDashboardOut,
    ThreeDDayStatusUpdate,
    ThreeDEdgeCommandAck,
    ThreeDEdgeCommandClaim,
    ThreeDEdgeHeartbeat,
    ThreeDEdgeStatusBatch,
    ThreeDInventoryAdjustment,
    ThreeDInventoryMovementOut,
    ThreeDInventoryOut,
    ThreeDMaintenanceInput,
    ThreeDMaintenanceOut,
    ThreeDMaintenanceUpdate,
    ThreeDMaterialInput,
    ThreeDMaterialOut,
    ThreeDMaterialUpdate,
    ThreeDPrinterCommandCreate,
    ThreeDPrinterCommandOut,
    ThreeDProductInput,
    ThreeDProductionRecordInput,
    ThreeDProductionRecordOut,
    ThreeDProductionRecordUpdate,
    ThreeDProductOut,
    ThreeDProductUpdate,
    ThreeDScheduleInput,
    ThreeDScheduleOut,
    ThreeDScheduleStatusUpdate,
    ThreeDScheduleUpdate,
    ThreeDSettingsOut,
    ThreeDSettingsUpdate,
    ThreeDStockInInput,
)
from app.services.auth import (
    AuthContext,
    ensure_permission_in_scope,
    get_current_user,
    has_permission_in_scope,
)
from app.services.three_d_migration_queries import (
    get_migration_batch,
    get_migration_reconciliation,
    list_migration_batches,
    list_migration_rows,
    require_migration_administrator,
)
from app.services.three_d_printing import (
    MAX_PRODUCT_IMAGE_BYTES,
    THREE_D_DEPARTMENTS,
    acknowledge_printer_command,
    adjust_inventory,
    apply_edge_status_batch,
    archive_material,
    archive_product,
    claim_printer_commands,
    command_out,
    create_maintenance,
    create_material,
    create_printer_command,
    create_product,
    create_production_record,
    create_schedule,
    create_stock_in,
    dashboard_snapshot,
    delete_maintenance,
    delete_production_record,
    delete_schedule,
    inventory_movement_out,
    inventory_out,
    list_audit_events,
    maintenance_out,
    material_out,
    product_image_path,
    product_out,
    production_record_out,
    register_edge_agent,
    remove_product_image,
    require_three_d_factory,
    save_product_image,
    schedule_out,
    update_day_status,
    update_maintenance,
    update_material,
    update_product,
    update_production_record,
    update_schedule,
    update_schedule_status,
    update_settings,
)

router = APIRouter(prefix="/api/three-d-printing", tags=["three-d-printing"])
MigrationDb = Annotated[Session, Depends(get_db)]
MigrationUser = Annotated[AuthContext, Depends(get_current_user)]


@router.get("/migration-batches")
def get_migration_batches(
    factory_id: str,
    db: MigrationDb,
    current_user: MigrationUser,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=100),
):
    factory_id = require_migration_administrator(current_user, factory_id)
    return list_migration_batches(db, factory_id=factory_id, page=page, page_size=page_size)


@router.get("/migration-batches/{batch_id}")
def get_migration_batch_detail(
    batch_id: str,
    factory_id: str,
    db: MigrationDb,
    current_user: MigrationUser,
):
    factory_id = require_migration_administrator(current_user, factory_id)
    return get_migration_batch(db, factory_id=factory_id, batch_id=batch_id)


@router.get("/migration-batches/{batch_id}/rows")
def get_migration_batch_rows(
    batch_id: str,
    factory_id: str,
    db: MigrationDb,
    current_user: MigrationUser,
    status: str = "",
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=100),
):
    factory_id = require_migration_administrator(current_user, factory_id)
    return list_migration_rows(db, factory_id=factory_id, batch_id=batch_id, status=status, page=page, page_size=page_size)


@router.get("/migration-batches/{batch_id}/row-errors")
def get_migration_batch_row_errors(
    batch_id: str,
    factory_id: str,
    db: MigrationDb,
    current_user: MigrationUser,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=100),
):
    factory_id = require_migration_administrator(current_user, factory_id)
    return list_migration_rows(db, factory_id=factory_id, batch_id=batch_id, status="failed", page=page, page_size=page_size)


@router.get("/migration-batches/{batch_id}/reconciliation")
def get_migration_batch_reconciliation(
    batch_id: str,
    factory_id: str,
    db: MigrationDb,
    current_user: MigrationUser,
):
    factory_id = require_migration_administrator(current_user, factory_id)
    return get_migration_reconciliation(db, factory_id=factory_id, batch_id=batch_id)


def _request_id(request: Request) -> str:
    return request.headers.get("x-request-id", "").strip()


def _ensure_permission(
    db: Session,
    user: AuthContext,
    permission: str,
    factory_id: str,
) -> None:
    factory_id = require_three_d_factory(factory_id)
    if any(
        has_permission_in_scope(user, permission, factory_id, department)
        for department in THREE_D_DEPARTMENTS
    ):
        return
    ensure_permission_in_scope(
        db,
        user,
        permission,
        factory_id,
        THREE_D_DEPARTMENTS[0],
    )


def _ensure_edge_token(x_edge_token: str) -> None:
    configured = settings.three_d_edge_agent_token.strip()
    supplied = x_edge_token.strip()
    if not configured:
        raise HTTPException(status_code=503, detail="云端尚未配置3D边缘代理令牌")
    if not supplied or not compare_digest(configured, supplied):
        raise HTTPException(status_code=401, detail="3D边缘代理认证失败")


@router.get("/dashboard", response_model=ThreeDDashboardOut)
def get_dashboard(
    factory_id: str,
    date_from: str = "",
    date_to: str = "",
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:read", factory_id)
    return dashboard_snapshot(
        db,
        factory_id=factory_id,
        date_from=date_from,
        date_to=date_to,
    )


@router.put("/settings", response_model=ThreeDSettingsOut)
def put_settings(
    payload: ThreeDSettingsUpdate,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", payload.factory_id)
    return update_settings(db, payload, current_user, _request_id(request))


@router.post("/materials", response_model=ThreeDMaterialOut, status_code=201)
def post_material(
    payload: ThreeDMaterialInput,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", payload.factory_id)
    return material_out(create_material(db, payload, current_user, _request_id(request)))


@router.put("/materials/{material_id}", response_model=ThreeDMaterialOut)
def put_material(
    material_id: str,
    payload: ThreeDMaterialUpdate,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", payload.factory_id)
    return material_out(
        update_material(db, material_id, payload, current_user, _request_id(request))
    )


@router.delete("/materials/{material_id}", status_code=204)
def delete_material(
    material_id: str,
    factory_id: str,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", factory_id)
    archive_material(db, material_id, factory_id, current_user, _request_id(request))


@router.post("/products", response_model=ThreeDProductOut, status_code=201)
def post_product(
    payload: ThreeDProductInput,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", payload.factory_id)
    return product_out(create_product(db, payload, current_user, _request_id(request)))


@router.put("/products/{product_id}", response_model=ThreeDProductOut)
def put_product(
    product_id: str,
    payload: ThreeDProductUpdate,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", payload.factory_id)
    return product_out(
        update_product(db, product_id, payload, current_user, _request_id(request))
    )


@router.delete("/products/{product_id}", status_code=204)
def delete_product(
    product_id: str,
    factory_id: str,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", factory_id)
    archive_product(db, product_id, factory_id, current_user, _request_id(request))


@router.post("/products/{product_id}/image", response_model=ThreeDProductOut)
async def post_product_image(
    product_id: str,
    request: Request,
    factory_id: str = Query(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:image_upload", factory_id)
    content = await file.read(MAX_PRODUCT_IMAGE_BYTES + 1)
    image = save_product_image(
        db,
        product_id=product_id,
        factory_id=factory_id,
        file_name=file.filename or "",
        mime_type=(file.content_type or "").lower(),
        content=content,
        user=current_user,
        request_id=_request_id(request),
    )
    product = db.get(ThreeDPrintingProduct, product_id)
    return product_out(product, image)


@router.get("/products/{product_id}/image")
def get_product_image(
    product_id: str,
    factory_id: str,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:read", factory_id)
    image = db.scalar(
        select(ThreeDPrintingProductImage).where(
            ThreeDPrintingProductImage.product_id == product_id,
            ThreeDPrintingProductImage.factory_id == factory_id,
            ThreeDPrintingProductImage.is_current.is_(True),
        )
    )
    if image is None:
        raise HTTPException(status_code=404, detail="产品图片不存在")
    path = product_image_path(image)
    if not path.is_file():
        raise HTTPException(status_code=404, detail="产品图片文件缺失")
    return FileResponse(
        path,
        media_type=image.mime_type,
        headers={
            "Cache-Control": "private, max-age=3600",
            "ETag": f'"{image.sha256}"',
        },
    )


@router.delete("/products/{product_id}/image", status_code=204)
def delete_product_image(
    product_id: str,
    factory_id: str,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:image_upload", factory_id)
    remove_product_image(
        db,
        product_id=product_id,
        factory_id=factory_id,
        user=current_user,
        request_id=_request_id(request),
    )


@router.post(
    "/records",
    response_model=ThreeDProductionRecordOut,
    status_code=201,
)
def post_record(
    payload: ThreeDProductionRecordInput,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", payload.factory_id)
    return production_record_out(
        create_production_record(db, payload, current_user, _request_id(request))
    )


@router.put("/records/{record_id}", response_model=ThreeDProductionRecordOut)
def put_record(
    record_id: str,
    payload: ThreeDProductionRecordUpdate,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", payload.factory_id)
    return production_record_out(
        update_production_record(
            db,
            record_id,
            payload,
            current_user,
            _request_id(request),
        )
    )


@router.delete("/records/{record_id}", status_code=204)
def delete_record(
    record_id: str,
    factory_id: str,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", factory_id)
    delete_production_record(
        db,
        record_id,
        factory_id,
        current_user,
        _request_id(request),
    )


@router.put("/day-status", status_code=200)
def put_day_status(
    payload: ThreeDDayStatusUpdate,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", payload.factory_id)
    record = update_day_status(db, payload, current_user, _request_id(request))
    return {
        "factory_id": record.factory_id,
        "business_date": record.business_date,
        "is_day_off": record.is_day_off,
        "revision": record.revision,
    }


@router.post("/inventory/adjust", response_model=ThreeDInventoryOut)
def post_inventory_adjustment(
    payload: ThreeDInventoryAdjustment,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", payload.factory_id)
    return inventory_out(
        adjust_inventory(db, payload, current_user, _request_id(request))
    )


@router.post(
    "/inventory/stock-in",
    response_model=ThreeDInventoryMovementOut,
    status_code=201,
)
def post_stock_in(
    payload: ThreeDStockInInput,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", payload.factory_id)
    return inventory_movement_out(
        create_stock_in(db, payload, current_user, _request_id(request))
    )


@router.post("/schedules", response_model=ThreeDScheduleOut, status_code=201)
def post_schedule(
    payload: ThreeDScheduleInput,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", payload.factory_id)
    return schedule_out(create_schedule(db, payload, current_user, _request_id(request)))


@router.put("/schedules/{schedule_id}", response_model=ThreeDScheduleOut)
def put_schedule(
    schedule_id: str,
    payload: ThreeDScheduleUpdate,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", payload.factory_id)
    return schedule_out(
        update_schedule(
            db,
            schedule_id,
            payload,
            current_user,
            _request_id(request),
        )
    )


@router.put("/schedules/{schedule_id}/status", response_model=ThreeDScheduleOut)
def put_schedule_status(
    schedule_id: str,
    payload: ThreeDScheduleStatusUpdate,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", payload.factory_id)
    return schedule_out(
        update_schedule_status(
            db,
            schedule_id,
            payload,
            current_user,
            _request_id(request),
        )
    )


@router.delete("/schedules/{schedule_id}", status_code=204)
def remove_schedule(
    schedule_id: str,
    factory_id: str,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", factory_id)
    delete_schedule(
        db,
        schedule_id,
        factory_id,
        current_user,
        _request_id(request),
    )


@router.post(
    "/maintenance",
    response_model=ThreeDMaintenanceOut,
    status_code=201,
)
def post_maintenance(
    payload: ThreeDMaintenanceInput,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", payload.factory_id)
    return maintenance_out(
        create_maintenance(db, payload, current_user, _request_id(request))
    )


@router.put("/maintenance/{maintenance_id}", response_model=ThreeDMaintenanceOut)
def put_maintenance(
    maintenance_id: str,
    payload: ThreeDMaintenanceUpdate,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", payload.factory_id)
    return maintenance_out(
        update_maintenance(
            db,
            maintenance_id,
            payload,
            current_user,
            _request_id(request),
        )
    )


@router.delete("/maintenance/{maintenance_id}", status_code=204)
def remove_maintenance(
    maintenance_id: str,
    factory_id: str,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:operate", factory_id)
    delete_maintenance(
        db,
        maintenance_id,
        factory_id,
        current_user,
        _request_id(request),
    )


@router.post(
    "/printers/{printer_id}/commands",
    response_model=ThreeDPrinterCommandOut,
    status_code=202,
)
def post_printer_command(
    printer_id: str,
    payload: ThreeDPrinterCommandCreate,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(
        db,
        current_user,
        "three_d_printing:printer_control",
        payload.factory_id,
    )
    return command_out(
        create_printer_command(
            db,
            printer_id=printer_id,
            payload=payload,
            user=current_user,
            request_id=_request_id(request),
        )
    )


@router.get("/audit", response_model=list[ThreeDAuditEventOut])
def get_audit(
    factory_id: str,
    limit: int = Query(default=200, ge=1, le=1000),
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:audit_read", factory_id)
    return list_audit_events(db, factory_id=factory_id, limit=limit)


@router.get("/export.xlsx")
def export_workbook(
    factory_id: str,
    date_from: str = "",
    date_to: str = "",
    db: Session = Depends(get_db),
    current_user: AuthContext = Depends(get_current_user),
):
    _ensure_permission(db, current_user, "three_d_printing:export", factory_id)
    snapshot = dashboard_snapshot(
        db,
        factory_id=factory_id,
        date_from=date_from,
        date_to=date_to,
    )
    workbook = Workbook()
    record_sheet = workbook.active
    record_sheet.title = "每日生产记录"
    sheets = [
        (
            record_sheet,
            [
                "日期",
                "机台",
                "状态",
                "产品",
                "客户",
                "材料",
                "料重(g)",
                "数量",
                "耗时(h)",
                "设计费",
                "报价",
                "备注",
                "开始时间",
                "完成时间",
            ],
            [
                [
                    row["business_date"],
                    row["machine_no"],
                    row["status"],
                    row["product_name"],
                    row["customer"],
                    row["material_name"],
                    row["weight_g"],
                    row["quantity"],
                    row["duration_hours"],
                    row["design_fee"],
                    row["quoted_price"],
                    row["remark"],
                    row["print_start_at"],
                    row["print_end_at"],
                ]
                for row in snapshot["records"]
            ],
        ),
        (
            workbook.create_sheet("产品库"),
            ["产品", "客户", "材料", "料重(g)", "耗时(h)", "数量", "默认报价", "图片SHA256"],
            [
                [
                    row["name"],
                    row["customer"],
                    row["material_name"],
                    row["weight_g"],
                    row["duration_hours"],
                    row["default_quantity"],
                    row["quoted_price"],
                    row["image_sha256"],
                ]
                for row in snapshot["products"]
            ],
        ),
        (
            workbook.create_sheet("库存"),
            ["材料", "当前库存(g)", "警戒库存(g)", "状态"],
            [
                [
                    row["material_name"],
                    row["stock_g"],
                    row["min_stock_g"],
                    "低库存" if row["is_low"] else "正常",
                ]
                for row in snapshot["inventory"]
            ],
        ),
        (
            workbook.create_sheet("排期表"),
            ["日期", "产品", "客户", "材料", "料重(g)", "数量", "机台", "优先级", "状态", "备注"],
            [
                [
                    row["business_date"],
                    row["product_name"],
                    row["customer"],
                    row["material_name"],
                    row["weight_g"],
                    row["quantity"],
                    row["machine_no"],
                    row["priority"],
                    row["status"],
                    row["remark"],
                ]
                for row in snapshot["schedules"]
            ],
        ),
        (
            workbook.create_sheet("维修记录"),
            ["日期", "机台", "类型", "描述", "费用", "服务商", "备注"],
            [
                [
                    row["business_date"],
                    row["machine_no"],
                    row["maintenance_type"],
                    row["description"],
                    row["cost"],
                    row["vendor"],
                    row["remark"],
                ]
                for row in snapshot["maintenance"]
            ],
        ),
    ]
    header_fill = PatternFill("solid", fgColor="0F766E")
    for sheet, headers, rows in sheets:
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_cells in sheet.columns:
            width = min(
                max(len(str(cell.value or "")) for cell in column_cells) + 2,
                36,
            )
            sheet.column_dimensions[column_cells[0].column_letter].width = max(width, 10)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                'attachment; filename="three-d-printing-huakang-a.xlsx"'
            )
        },
    )


@router.post("/edge/heartbeat")
def edge_heartbeat(
    payload: ThreeDEdgeHeartbeat,
    x_edge_token: str = Header(default=""),
    db: Session = Depends(get_db),
):
    _ensure_edge_token(x_edge_token)
    agent = register_edge_agent(db, payload)
    return {
        "agent_id": agent.id,
        "factory_id": agent.factory_id,
        "server_time": agent.last_seen_at,
        "command_poll_interval_seconds": settings.three_d_command_poll_interval_seconds,
    }


@router.post("/edge/status")
def edge_status(
    payload: ThreeDEdgeStatusBatch,
    x_edge_token: str = Header(default=""),
    db: Session = Depends(get_db),
):
    _ensure_edge_token(x_edge_token)
    printers = apply_edge_status_batch(
        db,
        factory_id=payload.factory_id,
        agent_key=payload.agent_key,
        statuses=payload.statuses,
    )
    return {
        "ok": True,
        "accepted": len(printers),
        "server_time": datetime.now(UTC).isoformat(),
    }


@router.post("/edge/commands/claim")
def edge_claim_commands(
    payload: ThreeDEdgeCommandClaim,
    x_edge_token: str = Header(default=""),
    db: Session = Depends(get_db),
):
    _ensure_edge_token(x_edge_token)
    commands = claim_printer_commands(
        db,
        factory_id=payload.factory_id,
        agent_key=payload.agent_key,
        limit=payload.limit,
    )
    printers = {
        printer.id: printer
        for printer in db.scalars(
            select(ThreeDPrintingPrinter).where(
                ThreeDPrintingPrinter.factory_id == payload.factory_id
            )
        )
    }
    return {
        "commands": [
            {
                **command_out(command),
                "machine_no": printers[command.printer_id].machine_no,
                "printer_type": printers[command.printer_id].printer_type,
            }
            for command in commands
            if command.printer_id in printers
        ]
    }


@router.post("/edge/commands/{command_id}/ack")
def edge_ack_command(
    command_id: str,
    payload: ThreeDEdgeCommandAck,
    x_edge_token: str = Header(default=""),
    db: Session = Depends(get_db),
):
    _ensure_edge_token(x_edge_token)
    return command_out(
        acknowledge_printer_command(
            db,
            command_id=command_id,
            factory_id=payload.factory_id,
            agent_key=payload.agent_key,
            status=payload.status,
            message=payload.message,
        )
    )
