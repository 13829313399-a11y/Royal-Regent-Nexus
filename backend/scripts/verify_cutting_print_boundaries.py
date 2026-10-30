"""Generate and read back synthetic print files without opening any database.

Native Excel/PDF and physical-printer acceptance remain separate. Outputs are
exclusive directories under the repository's ignored .tmp tree only.
"""
from pathlib import Path
from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4
import argparse
import hashlib
import json
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT/'backend'))
sys.path.insert(0, str(ROOT/'backend/tests'))
from cutting_print_cases import print_cases
from app.services import cutting_export_workbook as workbooks
from openpyxl import load_workbook


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def generate(output):
    output = output.resolve()
    if not output.is_relative_to((ROOT/'.tmp').resolve()):
        raise ValueError('Print evidence must use a new directory inside repository .tmp')
    output.mkdir(parents=True, exist_ok=False)
    paths = [Path(__file__), ROOT/'backend/tests/cutting_print_cases.py',
             ROOT/'backend/app/services/cutting_export_workbook.py']
    source = {str(p.relative_to(ROOT)):digest(p) for p in paths}
    results = []
    for label, report in print_cases().items():
        original = json.dumps(report, ensure_ascii=False, sort_keys=True)
        payload, extension = workbooks.package(report, SimpleNamespace(id='synthetic-print', display_name='合成排版验收'),
            lambda value, user: workbooks.render_workbook(value, user, '合成打印验收样本'))
        assert extension == 'xlsx'
        assert original == json.dumps(report, ensure_ascii=False, sort_keys=True)
        book = load_workbook(BytesIO(payload), read_only=False)
        try:
            for table in report['sheets']:
                assert [list(r) for r in list(book[table['name']].values)[2:]] == table['rows']
            expected_cover = [[None if cell == '' else cell for cell in row]
                              for row in workbooks.print_summary(report)['rows']]
            assert [list(r) for r in list(book['打印摘要'].values)[2:]] == expected_cover
            assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)
            assert all(sheet.print_title_rows=='$1:$2' and str(sheet.page_setup.paperSize)==sheet.PAPERSIZE_A3
                       and sheet.page_setup.orientation=='landscape' for sheet in book)
            (output/(label+'.xlsx')).write_bytes(payload)
            (output/(label+'.source.json')).write_text(original, encoding='utf-8')
            results.append(dict(case=label, file=label+'.xlsx', sha256=hashlib.sha256(payload).hexdigest(),
                                sheets=len(book.sheetnames), source_rows=sum(len(t['rows']) for t in report['sheets'])))
        finally:
            book.close()
    assert source == {str(p.relative_to(ROOT)):digest(p) for p in paths}
    summary = dict(status='passed_xlsx_readback', cases=results, source=source, database_used=False,
                   native_excel_verified=False, physical_print_verified=False, release_ready=False)
    (output/'report.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    return summary


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path)
    args=parser.parse_args()
    output=args.output or ROOT/'.tmp/cutting-print-boundaries'/(
        datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')+'-'+uuid4().hex[:8])
    summary=generate(output)
    print(json.dumps(dict(output=str(output),status=summary['status'],cases=len(summary['cases']),database_used=False)))
    return 0


if __name__=='__main__':
    raise SystemExit(main())
