"""Isolated engine smoke test. Run on office-private, never against business files.

Requires ONLYOFFICE_JWT_SECRET, ONLYOFFICE_INTERNAL_URL (default onlyoffice),
and SMOKE_SOURCE_URL naming this container on the private network.
"""
import base64
import hashlib
import hmac
import json
import os
from io import BytesIO
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from threading import Thread
from uuid import uuid4

import httpx
from openpyxl import Workbook, load_workbook
from docx import Document
from pypdf import PdfReader


def main():
    secret = os.environ["ONLYOFFICE_JWT_SECRET"].encode()
    base = os.environ.get("ONLYOFFICE_INTERNAL_URL", "http://onlyoffice").rstrip("/")
    source = os.environ.get("SMOKE_SOURCE_URL", "http://rr-office-smoke:8080")
    files = {}
    book = Workbook(); book.active["A1"] = 7; book.active["A2"] = 8; book.active["A3"] = "=SUM(A1:A2)"
    stream = BytesIO(); book.save(stream); files["/formula.xlsx"] = stream.getvalue()
    word = Document(); word.add_paragraph("ONLYOFFICE private document smoke 20261026")
    stream = BytesIO(); word.save(stream); files["/document.docx"] = stream.getvalue()

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            data = files.get(self.path)
            if data is None:
                self.send_error(404); return
            self.send_response(200); self.send_header("Content-Length", str(len(data))); self.end_headers(); self.wfile.write(data)
        def log_message(self, *_args):
            pass

    server = ThreadingHTTPServer(("0.0.0.0", 8080), Handler)
    Thread(target=server.serve_forever, daemon=True).start()
    def b64(data): return base64.urlsafe_b64encode(data).decode().rstrip("=")
    def sign(data):
        message = b64(b'{"alg":"HS256","typ":"JWT"}') + "." + b64(json.dumps(data).encode())
        return message + "." + b64(hmac.new(secret, message.encode(), hashlib.sha256).digest())
    result = {}
    try:
        with httpx.Client(timeout=120, trust_env=False, follow_redirects=False) as client:
            assert client.get(base + "/healthcheck").text == "true"
            for filename, output in (("formula.xlsx", "xlsx"), ("formula.xlsx", "pdf"), ("document.docx", "pdf")):
                payload = {"async": False, "filetype": filename.rsplit(".", 1)[1], "key": uuid4().hex,
                           "outputtype": output, "title": filename, "url": source + "/" + filename}
                denied = client.post(base + "/converter", json=payload).json()
                assert denied.get("error") is not None, "Unsigned conversion must fail"
                response = client.post(base + "/converter", json=payload | {"token": sign(payload)})
                response.raise_for_status(); body = response.json()
                assert body.get("endConvert") is True, str(body)
                converted = client.get(body["fileUrl"]); converted.raise_for_status()
                if output == "xlsx":
                    assert load_workbook(BytesIO(converted.content)).active["A3"].value == "=SUM(A1:A2)"
                    cached = load_workbook(BytesIO(converted.content), data_only=True).active["A3"].value
                    # Same-format conversion can return source bytes without
                    # recalculating. This is not an interactive editor test.
                    result["xlsx_formula_preserved"] = {"formula": "=SUM(A1:A2)", "cached_result": cached, "bytes": len(converted.content)}
                else:
                    pdf = PdfReader(BytesIO(converted.content))
                    extracted = pdf.pages[0].extract_text()
                    if filename.endswith("docx"):
                        assert "ONLYOFFICE private document smoke" in extracted
                        result["word_pdf"] = {"pages": len(pdf.pages), "bytes": len(converted.content)}
                    else:
                        assert len(pdf.pages) >= 1 and len(extracted) > 0
                        result["spreadsheet_pdf"] = {"pages": len(pdf.pages), "text": extracted, "bytes": len(converted.content)}
            result["unsigned_requests_rejected"] = True
            print(json.dumps(result), flush=True)
    finally:
        server.shutdown(); server.server_close()


if __name__ == "__main__":
    main()
